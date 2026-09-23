"""Gate 10: QA de la v1.0 (Excel y Word) y revisión funcional de términos críticos.

Escribe resultados/gate10/QA_GATE10.json. No modifica entregables.
"""
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.dom import minidom

import docx
import pandas as pd
from openpyxl import load_workbook

SUB = Path(__file__).resolve().parents[1]
RES = SUB / "resultados"
V091 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v0.9.1_GATE09.xlsx"
V10 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v1.0.xlsx"
DOCX = RES / "REQUERIMIENTO_BETTERSOFT_REPORTE_REGULATORIO_UPLUS_v1.0.docx"
NUEVAS = {"00_METADATOS_V1", "18_CONFIRMACION_GATE10", "19_PENDIENTES_RESIDUALES", "20_MATRIZ_SOLICITUD_BETTERSOFT", "21_CHANGELOG_V091_V10"}
FILAS_AGREGADAS = {"00_LEEME": 5, "00_ESTADO_FASES": 1}
CELDAS_PERMITIDAS = {("05_DICCIONARIO_BETTERSOFT", "ESTADO"), ("07_BRECHAS_BETTERSOFT", "PRIORIDAD")}
H1 = ["1. Objetivo y principio", "2. Alcance", "3. Entidades y extractos", "4. Matriz de solicitud", "5. Datos que deben conservar historia",
      "6. Datos sin transformar y catálogos", "7. Códigos SIES que U+ puede entregar", "8. Variables que U+ no debe calcular", "9. Reglas de extracción",
      "10. Validaciones mínimas del reporte", "11. Aclaraciones que Bettersoft debe devolver", "Anexo. Cambios solicitados en U+"]
INTERNOS = ["HIPOTESIS", "HIPÓTESIS", "ipótesis", "C09-", "CT-0", "CT-1", "Gate", "GATE", ".py", ".tsv", ".json", "codigo_gobernanza", "decisión interna",
            "Decisión interna", "ontradic", "PENDIENTE_DEMOSTRAR", "PUENTE", "scratchpad", "auditor", "Auditor", "Hoja1", "IMPLEMENTACION", "implementación",
            "OBSERVADO", "DEFAULT_TECNICO", "SUPUESTO", "local_restricted", "descartad", "error histórico", "errores", "INCONSISTENCIA", "rectific"]
# Reglas descartadas en Gate 09: no deben aparecer como mapeo vigente
DESCARTADAS = {
    "TITULADO -> VIG": r"TITULAD\w*\s*(?:->|→)\s*(?:VIG\s*)?[02]",
    "SITUACION 49/27 -> FOR 3": r"(?:49|27)\b[^.;\n]{0,40}(?:->|→)\s*(?:FOR[_ A-Z]*)?3\b",
    "CONTINUIDAD -> FOR 2": r"CONTINUIDAD[^.;\n]{0,40}(?:->|→)\s*(?:FOR[_ A-Z]*)?2\b",
    "SIT_FON_SOL = 1": r"SIT_FON_SOL[^.;\n]{0,20}(?:=|->|→)\s*1\b",
    "S -> NB/X": r"\bS\s*(?:->|→)\s*(?:NB|X)\b",
    "O -> 4": r"\bO\s*(?:->|→)\s*4\b",
    "FOR 2 => 1900": r"FOR\s*2\s*=>\s*1900",
    "Periodo 3 -> 2": r"\b3\s*(?:->|→)\s*2\b",
    "Por definir -> 38": r"Por definir\s*(?:->|→)\s*38",
}
TERMINOS = ["VIGENCIA", "TITULADO", "CAMBIO INTERNO", "CAMBIO DE JORNADA", "CONTINUIDAD", "FORMA DE INGRESO", "FONDO SOLIDARIO",
            "PAÍS ESTUDIOS SECUNDARIOS", "JORNADA O", "SEXO S", "PERIODO 3", "CODIGO_UNICO", "VERSION"]
RUT_RE = re.compile(r"(?<![\w.])\d{1,2}\.?\d{3}\.?\d{3}-[\dkK](?![\w])|(?<![\w])\d{7,8}(?![\w])")


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def norm(t):
    return re.sub(r"\s+", " ", t.upper().replace("Í", "I").replace("É", "E").replace("Ó", "O").replace("Á", "A").replace("Ú", "U")
                  .replace("PERÍODO", "PERIODO").replace("_", " ").replace(" = ", " "))


def pii(texto):
    out = []
    for m in RUT_RE.finditer(texto):
        d = re.sub(r"\D", "", m.group(0).split("-")[0])
        if set(d) == {"9"} or d in {"12345678", "123456789"} or texto[:m.start()].rstrip().endswith("Commit"):
            continue
        out.append(m.group(0))
    return out


def qa_xlsx():
    r = {}
    wb1, wb0 = load_workbook(V10), load_workbook(V091)
    r["abre"] = True
    r["hojas"] = len(wb1.sheetnames)
    r["hojas_v091_conservadas"] = all(s in wb1.sheetnames for s in wb0.sheetnames)
    r["hojas_nuevas"] = sorted(set(wb1.sheetnames) - set(wb0.sheetnames))
    r["hojas_nuevas_esperadas"] = set(r["hojas_nuevas"]) == NUEVAS
    r["formulas"] = sum(1 for ws in wb1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    perdida, cambios, filas, encab, filtros, paneles = [], [], {}, [], [], []
    for s in wb0.sheetnames:
        a, b = wb0[s], wb1[s]
        filas[s] = (a.max_row - 1, b.max_row - 1)
        if b.max_row - a.max_row != FILAS_AGREGADAS.get(s, 0):
            perdida.append(f"{s}: filas {a.max_row - 1} -> {b.max_row - 1}")
        ha, hb = [c.value for c in a[1]], [c.value for c in b[1]]
        if hb[: len(ha)] != ha:
            encab.append(s)
        for row in range(2, a.max_row + 1):
            for col in range(1, a.max_column + 1):
                va, vb = a.cell(row=row, column=col).value, b.cell(row=row, column=col).value
                if va != vb:
                    cambios.append((s, ha[col - 1], a.cell(row=row, column=1).value))
        if a.auto_filter.ref and not b.auto_filter.ref:
            filtros.append(s)
        if a.freeze_panes and b.freeze_panes != a.freeze_panes:
            paneles.append(s)
    r["filas_por_hoja_v091_v10"] = filas
    r["perdida_de_filas"] = perdida
    r["encabezados_alterados"] = encab
    r["filtros_perdidos"] = filtros
    r["paneles_alterados"] = paneles
    r["celdas_modificadas"] = [list(c) for c in cambios]
    r["celdas_modificadas_no_permitidas"] = [list(c) for c in cambios if (c[0], c[1]) not in CELDAS_PERMITIDAS]
    for s in NUEVAS:
        ws = wb1[s]
        if not ws.auto_filter.ref or ws.freeze_panes != "A2":
            filtros.append(f"{s} (nueva)")
    meta = dict(pd.read_excel(V10, "00_METADATOS_V1", dtype=str).values)
    r["metadatos"] = meta
    r["metadatos_ok"] = meta.get("VERSION") == "1.0" and meta.get("ESTADO") == "CONGELADO" and meta.get("GATE_ORIGEN") == "GATE09" and meta.get("GATE_CONGELACION") == "GATE10"
    d = pd.read_excel(V10, "05_DICCIONARIO_BETTERSOFT", dtype=str).fillna("")
    mz = pd.read_excel(V10, "20_MATRIZ_SOLICITUD_BETTERSOFT", dtype=str).fillna("")
    pe = pd.read_excel(V10, "19_PENDIENTES_RESIDUALES", dtype=str).fillna("")
    cf = pd.read_excel(V10, "18_CONFIRMACION_GATE10", dtype=str).fillna("")
    sol = d[d.ACCION_BETTERSOFT.str.startswith("SOLICITAR")]
    r["columnas_bettersoft"] = len(sol)
    r["prioridades"] = sol.PRIORIDAD.value_counts().to_dict()
    r["P1_PENDIENTES"] = int(((d.PRIORIDAD == "P1") & (d.P1_DEMOSTRADO != "SI")).sum())
    r["matriz_filas"] = len(mz)
    r["matriz_coincide_con_diccionario"] = set(mz.ID) == set(sol.ID) and len(mz) == len(sol)
    r["matriz_campos_obligatorios_vacios"] = int((mz[["Campo solicitado", "Descripción funcional", "Fuente/entidad U+", "Granularidad", "Histórico", "Formato", "Prioridad"]] == "").sum().sum())
    r["duplicados_id"] = int(d.ID.duplicated().sum()) + int(mz.ID.duplicated().sum())
    r["pendientes_residuales"] = pe[["ID", "CAMPO", "ESTADO"]].values.tolist()
    r["pendientes_sexo_periodo_ok"] = all(((pe.CAMPO == c) & (pe.ESTADO == "PENDIENTE_CATALOGO_OFICIAL_UPLUS") & pe.VINCULO.str.contains("UP-68")).any()
                                          for c in ["SEXO_U_PLUS = S", "PERIODO_U_PLUS = 3"])
    r["confirmaciones"] = dict(zip(cf.ID, cf.ESTADO))
    fugas = []
    for s, df in pd.read_excel(V10, sheet_name=None, dtype=str).items():
        df = df.drop(columns=[c for c in df.columns if c == "bytes"])
        fugas += [(s, v[:30]) for v in df.fillna("").astype(str).values.ravel() if pii(v)]
    r["posibles_datos_personales"] = fugas[:5]
    # reglas descartadas en contenido vigente
    m = pd.read_excel(V10, "04_MATRIZ_MAESTRA", dtype=str).fillna("")
    c = pd.read_excel(V10, "06_CATALOGOS_HOMOLOGACIONES", dtype=str).fillna("")
    vigente = (list(m.MAPEO_VIGENTE) + list(sol["TRANSFORMACIÓN"]) + list(mz.astype(str).values.ravel())
               + list(c[c.CLASIFICACION_GATE09.isin(["OFICIAL", "OBSERVADO"])].HOMOLOGA_A))
    r["reglas_descartadas_en_vigente"] = {k: [v[:80] for v in vigente if re.search(p, v)] for k, p in DESCARTADAS.items() if any(re.search(p, v) for v in vigente)}
    ok = (r["hojas_v091_conservadas"] and r["hojas_nuevas_esperadas"] and not perdida and not encab and not filtros and not paneles
          and not r["celdas_modificadas_no_permitidas"] and r["metadatos_ok"] and r["P1_PENDIENTES"] == 0 and r["matriz_coincide_con_diccionario"]
          and not r["matriz_campos_obligatorios_vacios"] and not r["duplicados_id"] and r["pendientes_sexo_periodo_ok"] and not fugas
          and not r["reglas_descartadas_en_vigente"] and all(v in ("CONFIRMADA", "APROBADO") for v in r["confirmaciones"].values()))
    r["QA_XLSX_V1"] = "APROBADO" if ok else "NO_APROBADO"
    return r


def qa_docx():
    r = {}
    with zipfile.ZipFile(DOCX) as z:
        r["zip_integro"] = z.testzip() is None
        partes = z.namelist()
        r["partes"] = {p: p in partes for p in ["[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels", "word/numbering.xml", "word/footer1.xml"]}
        for p in [p for p in partes if p.endswith(".xml") or p.endswith(".rels")]:
            minidom.parseString(z.read(p))
        r["xml_bien_formado"] = True
        xml = z.read("word/document.xml").decode("utf-8")
    r["tabla_contenido"] = "TOC" in xml
    d = docx.Document(DOCX)
    h1 = [p.text for p in d.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    r["secciones_h1"] = h1
    r["secciones_ok"] = h1 == H1
    r["tablas"] = len(d.tables)
    vacias, actual, contenido = [], None, False
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            st = el.find(".//{*}pStyle")
            tx = "".join(t.text or "" for t in el.iter("{*}t")).strip()
            if st is not None and st.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "Heading1":
                if actual and not contenido:
                    vacias.append(actual)
                actual, contenido = tx, False
            elif tx:
                contenido = True
        elif tag == "tbl":
            contenido = True
    if actual and not contenido:
        vacias.append(actual)
    r["secciones_vacias"] = vacias
    texto = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])
    r["terminos_internos"] = sorted({t for t in INTERNOS if t in texto})
    r["posibles_datos_personales"] = pii(texto)[:5]
    r["reglas_descartadas"] = {k: re.findall(p, texto)[:3] for k, p in DESCARTADAS.items() if re.search(p, texto)}
    req = json.loads((RES / "gate10/req_v1.json").read_text(encoding="utf-8"))
    p1 = sum(1 for m in req["matriz"] if m["Prioridad"] == "P1")
    r["conteo_coherente"] = f"Se solicitan {len(req['matriz'])} campos" in texto and f"{p1} de prioridad P1" in texto
    r["matriz_filas_word"] = len(d.tables[1].rows) - 1
    r["pendientes_como_aclaracion"] = "SEXO = S" in texto and "PERÍODO = 3" in texto
    nt = norm(texto)
    r["terminos_criticos_word"] = {t: nt.count(norm(t)) for t in TERMINOS}
    ok = (r["zip_integro"] and all(r["partes"].values()) and r["secciones_ok"] and not vacias and not r["terminos_internos"]
          and not r["posibles_datos_personales"] and not r["reglas_descartadas"] and r["conteo_coherente"]
          and r["matriz_filas_word"] == len(req["matriz"]) and r["pendientes_como_aclaracion"])
    r["QA_DOCX_ESTRUCTURAL_V1"] = "APROBADO" if ok else "NO_APROBADO"
    return r


def main():
    out = {"xlsx": qa_xlsx(), "docx": qa_docx(), "sha256": {p.name: sha(p) for p in (V10, DOCX)}}
    (RES / "gate10").mkdir(exist_ok=True)
    (RES / "gate10/QA_GATE10.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    return 0 if out["xlsx"]["QA_XLSX_V1"] == "APROBADO" and out["docx"]["QA_DOCX_ESTRUCTURAL_V1"] == "APROBADO" else 1


if __name__ == "__main__":
    sys.exit(main())
