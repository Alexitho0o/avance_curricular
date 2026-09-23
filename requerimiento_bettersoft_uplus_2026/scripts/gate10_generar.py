"""Gate 10: congelación v1.0 del Diccionario Maestro U+ -> SIES y del requerimiento Bettersoft.

Parte de la v0.9.1 (Gate 09), que no se modifica. No cambia reglas funcionales: agrega metadatos de
congelación, confirmaciones humanas, pendientes residuales, la matriz de solicitud final y el changelog,
y corrige las inconsistencias inter-artefactos detectadas en la Fase 10.2 según el resultado aprobado
del Gate 09 (ACCION_BETTERSOFT del diccionario).

Salidas:
- resultados/DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v1.0.xlsx
- resultados/gate10/req_v1.json (insumo del Word v1.0)
- CHANGELOG_COLUMNAS_BETTERSOFT_v0.9.1_A_v1.0.md
- PENDIENTE_EVALUAR_RECTIFICACION_MU2026_C09_02_C09_03.md
"""
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SUB = Path(__file__).resolve().parents[1]
RES = SUB / "resultados"
V091 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v0.9.1_GATE09.xlsx"
V10 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v1.0.xlsx"
M = json.loads((RES / "gate09/METRICAS_EVIDENCIA_GATE09.json").read_text(encoding="utf-8"))
FECHA = "2026-09-23"
CONF = "Confirmación del responsable del subproyecto registrada en sesión de trabajo del 23-09-2026 (Gate humano A)"
PEND_CAT = "PENDIENTE_CATALOGO_OFICIAL_UPLUS"

METADATOS = [("VERSION", "1.0"), ("ESTADO", "CONGELADO"), ("FECHA_CONGELACION", FECHA), ("GATE_ORIGEN", "GATE09"),
             ("GATE_CONGELACION", "GATE10"), ("BASE", V091.name),
             ("PRINCIPIO", "Bettersoft entrega hechos y atributos del ERP; la institución gobierna la transformación regulatoria SIES.")]

CONFIRMACIONES = [
    ("C09-01", "TITULADO_NO_DETERMINA_VIGENCIA_SIES", "CONFIRMADA"),
    ("C09-02", "CAMBIO_JORNADA_NO_DETERMINA_FORMA_INGRESO_3", "CONFIRMADA"),
    ("C09-03", "NOMBRE_CONTINUIDAD_NO_DETERMINA_FORMA_INGRESO_2", "CONFIRMADA"),
    ("C09-04", "FONDO_SOLIDARIO_IPSS=0 / NO_CALCULAR_OTRO_VALOR_SIN_EVIDENCIA", "CONFIRMADA"),
    ("QA_DOCX_VISUAL v0.9.1", "Revisión manual en Microsoft Word de REQUERIMIENTO_..._v0.9.1_GATE09.docx", "APROBADO"),
    ("QA_DOCX_VISUAL v1.0", "Revisión manual en Microsoft Word de REQUERIMIENTO_BETTERSOFT_REPORTE_REGULATORIO_UPLUS_v1.0.docx (18 puntos)", "APROBADO"),
]

f3 = M["ct02"]["for3_2026_vs_mu2025"]
f2 = M["ct03"]["for2_2026_vs_mu2025"]
PENDIENTES = [
    dict(ID="PR-01", TIPO="Bettersoft", CAMPO="SEXO_U_PLUS = S", ESTADO=PEND_CAT, VINCULO="UP-10; UP-68; C09-05",
         DETALLE=f"Significado no documentado. {M['ct04']['uplus_S_filas']} filas U+ ({M['ct04']['uplus_S_rut']} RUT). No se infiere significado; no se asigna código SIES.",
         ACCION="Bettersoft documenta el significado en el catálogo U+ (UP-68). La institución decide luego la transformación por proceso.", BLOQUEA_P1="NO"),
    dict(ID="PR-02", TIPO="Bettersoft", CAMPO="PERIODO_U_PLUS = 3", ESTADO=PEND_CAT, VINCULO="UP-32; UP-40; UP-68; C09-09",
         DETALLE=f"Significado no documentado (PERIODOMATRICULA {M['uplus_0508']['PERIODOMATRICULA'].get('3', 0)} filas en el snapshot 08-05-2026). No se infiere semestre.",
         ACCION="Bettersoft documenta el catálogo de períodos con fechas de inicio y término por régimen (UP-68).", BLOQUEA_P1="NO"),
    dict(ID="PR-03", TIPO="Institucional (fuera del requerimiento Bettersoft)", CAMPO="PENDIENTE_EVALUAR_RECTIFICACION_MU2026_C09_02_C09_03",
         ESTADO="PENDIENTE_DECISION_INSTITUCIONAL", VINCULO="C09-02; C09-03",
         DETALLE=f"MU 2026 cargada: {M['ct02']['pes_for3']} filas FOR 3 ({f3.get('mismo_COD_CAR_distinta_JOR', 0) + f3.get('mismo_COD_CAR_misma_JOR', 0)} con el mismo COD_CAR en MU 2025) y "
                 f"{M['ct03']['pes_for2']} filas FOR 2 ({f2.get('for_2025_11', 0)} con FOR 11 en 2025 en el mismo programa).",
         ACCION="Evaluar si se solicita rectificación a SIES. No se rectifica ni se preparan cargas en este gate.", BLOQUEA_P1="NO"),
    dict(ID="PR-04", TIPO="Institucional (seguimiento)", CAMPO="Filas de Extranjeros Intercambio en la matriz", ESTADO="PENDIENTE_REVISION",
         VINCULO="MM-116 a MM-131", DETALLE="Contrastar con el Anexo II del Instructivo Extranjeros 2026, ya incorporado.", ACCION="Revisión documental; sin impacto en Bettersoft.", BLOQUEA_P1="NO"),
    dict(ID="PR-05", TIPO="Fuente", CAMPO="Manuales MU 2022 y 2023", ESTADO="FUENTE_NO_DISPONIBLE", VINCULO="C09-13",
         DETALLE="Solo se requieren para reconstrucción histórica.", ACCION="Incorporar si se ubican.", BLOQUEA_P1="NO"),
]

# Inconsistencias inter-artefactos (Fase 10.2): el anexo del Word v0.9.1 contradecía ACCION_BETTERSOFT del diccionario Gate 09
BR_V1 = {
    "BR-01": ("Almacenar el código único SIES de cada programa como dato maestro con historia, junto con sus atributos de oferta (UP-26, UP-27).", None, "Redacción para proveedor"),
    "BR-02": ("Registrar los hechos del ingreso a cada programa: vía de admisión y cupo, tipo de programa de origen, reconocimiento de acceso y trayectoria (UP-63 a UP-65, UP-34 a UP-37).", None, "Redacción para proveedor"),
    "BR-05": ("Entregar el detalle por asignatura con CODCLI (UP-62); los resúmenes los calcula la institución.", None, "Redacción para proveedor"),
    "BR-06": ("Entregar el detalle por asignatura (UP-62) y el catálogo de planes (UP-58); el avance lo calcula la institución.", None, "Redacción para proveedor"),
    "BR-07": ("Exponer el tipo de documento registrado en U+ (UP-03) y separar número y dígito verificador.", None,
              "INCONSISTENCIA_INTER_ARTEFACTOS: pedía el código SIES R/P, que el diccionario Gate 09 marca NO_CALCULAR (UP-04)"),
    "BR-10": ("Registrar la nacionalidad con un catálogo de países en U+ y entregar el valor registrado (UP-14).", None,
              "INCONSISTENCIA_INTER_ARTEFACTOS: pedía el código SIES 1-197, que el diccionario Gate 09 marca NO_CALCULAR (UP-15)"),
    "BR-15": ("Entregar el nivel registrado en U+ a la fecha de corte y el régimen del plan (UP-41, UP-30); documentar el valor 20 en el catálogo (UP-68).", None,
              "INCONSISTENCIA_INTER_ARTEFACTOS: pedía el nivel en semestres, que Gate 09 asigna a la institución"),
    "BR-16": ("Entregar el historial de estado y situación con fechas (UP-43); suspensiones y reincorporación las calcula la institución.", None, "Redacción para proveedor"),
    "BR-20": ("Entregar estado civil y comuna tal como están en U+ (UP-19, UP-21) y la dirección desagregada (UP-20).", None,
              "INCONSISTENCIA_INTER_ARTEFACTOS: pedía códigos FCU, que Gate 09 dejó como homologación institucional"),
    "BR-22": ("Excluir del reporte regulatorio las columnas sin uso SIES: responsable financiero, datos laborales, tramo de renta, número de hijos, contactos de emergencia, ejecutivo comercial y arancel.", None,
              "Referencia a hoja interna reemplazada por la lista explícita"),
    "BR-23": ("Registrar la fecha de egreso y la fecha de titulación de cada matrícula (UP-59, UP-60).", "P2",
              "INCONSISTENCIA_INTER_ARTEFACTOS: la brecha seguía en P4 aunque Gate 09 subió UP-59 y UP-60 a P2"),
    "BR-25": ("Entregar un catálogo documentado del significado de cada código U+ usado en los extractos (UP-68), incluidos SEXO = S y PERÍODO = 3.", None, "Vincula los pendientes residuales"),
}

ENTIDAD = {"PERSONA": "Ficha de la persona", "PERSONA-AÑO": "Ficha de la persona, por año", "MATRÍCULA": "Matrícula (CODCLI)",
           "MATRÍCULA-PERÍODO": "Matrícula, por año/período", "MATRÍCULA (historial)": "Historial de la matrícula",
           "MATRÍCULA-PERÍODO (historial)": "Historial de la matrícula, por año/período", "OFERTA": "Programa (maestro de oferta)",
           "MATRÍCULA-PERÍODO (atributo de OFERTA/PLAN)": "Programa (maestro de oferta), asignado a la matrícula por período",
           "PLAN": "Plan de estudios", "ASIGNATURA": "Registro académico por asignatura", "CATÁLOGO": "Tablas de parámetros de U+"}
GRANO = {"PERSONA": "Una fila por persona", "PERSONA-AÑO": "Una fila por persona y año", "MATRÍCULA": "Una fila por CODCLI",
         "MATRÍCULA-PERÍODO": "Una fila por CODCLI y año/período", "MATRÍCULA (historial)": "Una fila por cambio de estado de cada CODCLI",
         "MATRÍCULA-PERÍODO (historial)": "Una fila por CODCLI y año/período", "OFERTA": "Una fila por programa y vigencia",
         "MATRÍCULA-PERÍODO (atributo de OFERTA/PLAN)": "Una fila por programa y vigencia; referenciado por CODCLI y período",
         "PLAN": "Una fila por plan de estudios", "ASIGNATURA": "Una fila por asignatura cursada o reconocida", "CATÁLOGO": "Una fila por código"}
CATALOGO = {
    "UP-03": "Tipos de documento de U+", "UP-10": "Catálogo de SEXO de U+ (aclarar el código S)", "UP-14": "Sin catálogo hoy; valor tal como está en U+",
    "UP-16": "Tabla de países 1-197 (entregada por la institución) o nombre del país", "UP-17": "1 con residencia previa / 2 sin residencia previa / 3 no reside en Chile",
    "UP-18": "Tabla de países 1-197", "UP-19": "Catálogo de estado civil de U+", "UP-21": "Comunas de U+", "UP-23": "Catálogo de sedes de U+",
    "UP-24": "Catálogo de carreras de U+", "UP-25": "Catálogo de jornadas de U+ (aclarar el código O)",
    "UP-26": "Modalidad 1/2/3; tipo de plan regular / especial / regular de continuidad; nivel y duración según la oferta", "UP-27": "Códigos asignados por la institución",
    "UP-28": "Catálogo de planes de U+", "UP-29": "Correlativo por código único", "UP-30": "Catálogo de régimen de U+",
    "UP-32": "Catálogo de períodos de U+ con fechas (aclarar el período 3)", "UP-35": "Misma institución / otra institución / sin origen",
    "UP-40": "Catálogo de períodos de U+ con fechas (aclarar el período 3)", "UP-41": "Niveles de U+ (aclarar el valor 20)",
    "UP-42": "Catálogo de estados y situaciones de U+", "UP-43": "Catálogo de estados y situaciones de U+", "UP-58": "Tipo de unidad: asignaturas / créditos / otra",
    "UP-61": "Valores de MATRICULA de U+ (aclarar 1 y 2)", "UP-62": "Tipo de registro: cursada / convalidada / homologada / reconocida",
    "UP-63": "Vías de admisión de U+ y tipos de cupo especial", "UP-64": "Plan común o bachillerato / técnico de nivel superior / profesional / otro / sin programa previo",
    "UP-65": "RAP / reconocimiento de estudios / ninguno", "UP-66": "Los mismos catálogos de sede, jornada, plan, nivel y estado", "UP-68": "—",
}
DISP = {"EXISTE": "Existe en U+", "NO_EXISTE": "No existe en U+: incorporar", "POR_CONFIRMAR": "Bettersoft debe confirmar si existe", "EXISTE PARCIAL": "Existe en parte"}

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
NEW_FILL = PatternFill("solid", fgColor="375623")
WHITE = Font(color="FFFFFF", bold=True)


def hmap(ws):
    return {c.value: c.column for c in ws[1] if c.value is not None}


def add_cols(ws, names):
    hm = hmap(ws)
    for n in names:
        if n not in hm:
            col = ws.max_column + 1
            c = ws.cell(row=1, column=col, value=n)
            c.fill, c.font = NEW_FILL, WHITE
            ws.column_dimensions[get_column_letter(col)].width = 34
            hm[n] = col
    return hm


def append(ws, hm, data):
    r = ws.max_row + 1
    for k, v in data.items():
        if k in hm:
            ws.cell(row=r, column=hm[k], value=v).alignment = Alignment(wrap_text=True, vertical="top")


def sheet(wb, name, cols, rows, widths=None, pos=None):
    ws = wb.create_sheet(name, pos)
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r.get(c, "") if isinstance(r, dict) else r[i] for i, c in enumerate(cols)])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 26)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def limpia_fuente(t):
    t = str(t or "")
    t = t.replace("Hoja1 (PROMEDIOSDEALUMNOS)", "reporte de notas 7804").replace("Hoja1 ", "reporte de notas 7804: ")
    return t.split(" - hoy se infiere")[0].replace(" (significado no documentado)", "")


def main():
    wb = load_workbook(V091)
    changelog = []

    # metadatos
    sheet(wb, "00_METADATOS_V1", ["CLAVE", "VALOR"], METADATOS, {"CLAVE": 24, "VALOR": 90}, pos=0)
    ws = wb["00_LEEME"]
    hm = hmap(ws)
    for k, v in METADATOS[:5]:
        append(ws, hm, {"TEMA": k, "DETALLE": v})

    # diccionario: pendientes residuales
    ws = wb["05_DICCIONARIO_BETTERSOFT"]
    hm = add_cols(ws, ["PENDIENTE_RESIDUAL_V1"])
    dic_rows = []
    for row in range(2, ws.max_row + 1):
        uid = ws.cell(row=row, column=hm["ID"]).value
        est = ws.cell(row=row, column=hm["ESTADO"]).value
        nuevo, pr = None, None
        if uid == "UP-10":
            nuevo, pr = f"OBSERVADO M/F; S = {PEND_CAT} (UP-68)", "PR-01"
        elif uid == "UP-32":
            nuevo, pr = f"1/2 OBSERVADO; 3 = {PEND_CAT} (UP-68)", "PR-02"
        elif uid == "UP-40":
            pr = "PR-02"
        elif uid == "UP-68":
            nuevo, pr = "PENDIENTE_DATO_UPLUS; incluye PR-01 (SEXO = S) y PR-02 (PERÍODO = 3)", "PR-01; PR-02"
        if pr:
            ws.cell(row=row, column=hm["PENDIENTE_RESIDUAL_V1"], value=pr)
        if nuevo:
            ws.cell(row=row, column=hm["ESTADO"], value=nuevo)
            changelog.append(dict(CAMPO=f"{uid} {ws.cell(row=row, column=hm['CAMPO SOLICITADO']).value}", ATRIBUTO="ESTADO", V091=est, V10=nuevo,
                                  CLASE="PENDIENTE_BETTERSOFT", MOTIVO="Fase 10.3: pendiente residual explícito vinculado a UP-68; no se infiere significado"))
        dic_rows.append({c: ws.cell(row=row, column=i).value for c, i in hm.items()})

    # brechas
    ws = wb["07_BRECHAS_BETTERSOFT"]
    hm = add_cols(ws, ["CAMBIO_SOLICITADO_V1", "NOTA_V1"])
    br_rows = []
    for row in range(2, ws.max_row + 1):
        bid = ws.cell(row=row, column=hm["ID"]).value
        texto, pri, nota = BR_V1.get(bid, (None, None, ""))
        base = ws.cell(row=row, column=hm["CAMBIO SOLICITADO"]).value
        g09 = ws.cell(row=row, column=hm["CAMBIO_GATE09"]).value
        reformulada = ws.cell(row=row, column=hm["ESTADO_GATE09"]).value in ("ACTUALIZADA", "REFORMULADA", "AMPLIADA")
        ws.cell(row=row, column=hm["CAMBIO_SOLICITADO_V1"], value=texto or (g09 if reformulada else base))
        ws.cell(row=row, column=hm["NOTA_V1"], value=nota)
        if pri:
            ant = ws.cell(row=row, column=hm["PRIORIDAD"]).value
            ws.cell(row=row, column=hm["PRIORIDAD"], value=pri)
            changelog.append(dict(CAMPO=f"{bid} (brecha)", ATRIBUTO="PRIORIDAD", V091=ant, V10=pri, CLASE="RECLASIFICADO",
                                  MOTIVO="Alineación con UP-59/UP-60 (P2 en Gate 09); no es una nueva decisión de prioridad"))
        if nota.startswith("INCONSISTENCIA"):
            changelog.append(dict(CAMPO=f"{bid} (brecha)", ATRIBUTO="CAMBIO SOLICITADO", V091=ws.cell(row=row, column=hm["CAMBIO_GATE09"]).value or base,
                                  V10=texto, CLASE="MODIFICADO", MOTIVO=nota))
        br_rows.append({c: ws.cell(row=row, column=i).value for c, i in hm.items()})

    # estado de fases
    ws = wb["00_ESTADO_FASES"]
    append(ws, hmap(ws), {"FASE": "10 - Gate 10 congelación v1.0", "ESTADO": "CONGELADO (ver MANIFIESTO_CONGELACION_v1.0.md)",
                          "PROCESOS_ANALIZADOS": "—", "ARCHIVOS_REVISADOS": "Artefactos Gate 09", "CAMPOS_IDENTIFICADOS": "Sin cambios de reglas",
                          "CRUCES_CONFIRMADOS": "C09-01 a C09-04 confirmados institucionalmente", "PENDIENTES": "PR-01 a PR-05",
                          "BLOQUEOS": "Ninguno", "SIGUIENTE_FASE": "Envío a Bettersoft"})

    # matriz de solicitud final (Fase 10.8)
    sol = [r for r in dic_rows if r["ACCION_BETTERSOFT"] in ("SOLICITAR_HECHO", "SOLICITAR_MAESTRO")]
    orden = {"P1": 1, "P2": 2, "P3": 3, "P4": 4}
    sol.sort(key=lambda r: (orden[r["PRIORIDAD"]], int(r["ID"][3:])))
    matriz = []
    for i, r in enumerate(sol, 1):
        obs = [DISP.get(r["DISPONIBILIDAD_UPLUS"], r["DISPONIBILIDAD_UPLUS"]), f"Validación: {r['VALIDACION']}" if r["VALIDACION"] not in ("", "-", "—", None) else ""]
        if r["ACCION_BETTERSOFT"] == "SOLICITAR_MAESTRO":
            obs.append("Dato maestro: lo asigna la institución y U+ lo almacena con vigencia")
        if r.get("DATO_SENSIBLE", "").startswith("SI"):
            obs.append("Dato personal o sensible")
        matriz.append({"Nº": i, "ID": r["ID"], "Campo solicitado": r["CAMPO SOLICITADO"], "Descripción funcional": r["DEFINICIÓN FUNCIONAL"],
                       "Fuente/entidad U+": f"{ENTIDAD.get(r['NIVEL'], r['NIVEL'])}. Hoy: {limpia_fuente(r['CAMPO U+ ACTUAL'])}",
                       "Granularidad": GRANO.get(r["NIVEL"], r["NIVEL"]), "Histórico": "SI" if r["REQUERIMIENTO_HISTORICO"] == "SI" else "NO",
                       "Formato": r["TIPO"] + ("; dd/mm/aaaa" if "Fecha" in (r["TIPO"] or "") else ""), "Catálogo": CATALOGO.get(r["ID"], "—"),
                       "Prioridad": r["PRIORIDAD"], "Procesos SIES": r["PROCESOS QUE LO UTILIZAN"], "Observaciones": "; ".join(o for o in obs if o)})
    cols_m = ["Nº", "ID", "Campo solicitado", "Descripción funcional", "Fuente/entidad U+", "Granularidad", "Histórico", "Formato", "Catálogo", "Prioridad", "Procesos SIES", "Observaciones"]
    sheet(wb, "20_MATRIZ_SOLICITUD_BETTERSOFT", cols_m, matriz, {"Campo solicitado": 32, "Descripción funcional": 60, "Fuente/entidad U+": 36, "Observaciones": 40})

    # confirmaciones y pendientes
    sheet(wb, "18_CONFIRMACION_GATE10", ["ID", "RESOLUCION", "ESTADO", "FECHA", "REGISTRO"], [(a, b, c, FECHA, CONF if a.startswith("C09") else "Revisión manual en Microsoft Word (Gate humano B)") for a, b, c in CONFIRMACIONES],
          {"RESOLUCION": 60, "REGISTRO": 60})
    sheet(wb, "19_PENDIENTES_RESIDUALES", ["ID", "TIPO", "CAMPO", "ESTADO", "VINCULO", "DETALLE", "ACCION", "BLOQUEA_P1"], PENDIENTES, {"DETALLE": 60, "ACCION": 50})

    # changelog de columnas Bettersoft (Fase 10.7)
    ids_cambio = {c["CAMPO"].split()[0] for c in changelog if c["CAMPO"].startswith("UP-")}
    sin_cambio = [r["ID"] for r in dic_rows if r["ID"] not in ids_cambio]
    cols_c = ["CAMPO", "ATRIBUTO", "V091", "V10", "CLASE", "MOTIVO"]
    sheet(wb, "21_CHANGELOG_V091_V10", cols_c, changelog + [dict(CAMPO=f"{len(sin_cambio)} filas del diccionario", ATRIBUTO="—", V091="—", V10="—", CLASE="SIN_CAMBIO",
                                                              MOTIVO=", ".join(sin_cambio))], {"V091": 45, "V10": 45, "MOTIVO": 70})
    wb.save(V10)

    # JSON para el Word v1.0
    br_word = [{"ID": b["ID"], "CAMPO": b["CAMPO"], "CAMBIO": b["CAMBIO_SOLICITADO_V1"], "PROCESOS": b["PROCESOS AFECTADOS"], "PRIORIDAD": b["PRIORIDAD"]} for b in br_rows]
    no_calc = json.loads((RES / "gate09/req_gate09.json").read_text(encoding="utf-8"))
    req = {"matriz": matriz, "br": br_word, "no_calcular": no_calc["no_calcular"], "codigo_unico": no_calc["codigo_unico"], "historial": no_calc["historial"],
           "dic": [{k: r.get(k) for k in ("ID", "CAMPO SOLICITADO", "EJEMPLO", "VALIDACION", "BLOQUE", "PRIORIDAD", "DISPONIBILIDAD_UPLUS")} for r in sol]}
    (RES / "gate10").mkdir(exist_ok=True)
    (RES / "gate10/req_v1.json").write_text(json.dumps(req, ensure_ascii=False, indent=1), encoding="utf-8")

    # changelog md
    pr = {p: sum(1 for r in sol if r["PRIORIDAD"] == p) for p in ("P1", "P2", "P3", "P4")}
    lines = ["# Changelog de columnas Bettersoft: v0.9.1 → v1.0", "",
             f"Fecha: {FECHA}. Base: `{V091.name}`. La v1.0 no cambia reglas funcionales ni prioridades de campos.", "",
             "## Conteos", "", "| Versión | Filas del diccionario | Solicitadas a Bettersoft | P1 | P2 | P3 | P4 | No calcular en U+ |", "|---|---|---|---|---|---|---|---|",
             "| v0.9 | 61 | 61 | 25 | 23 | 8 | 5 | 0 |",
             f"| v0.9.1 (Gate 09) | {len(dic_rows)} | {len(sol)} | {pr['P1']} | {pr['P2']} | {pr['P3']} | {pr['P4']} | {sum(1 for r in dic_rows if r['ACCION_BETTERSOFT'] == 'NO_CALCULAR_EN_UPLUS')} |",
             f"| v1.0 (Gate 10) | {len(dic_rows)} | {len(sol)} | {pr['P1']} | {pr['P2']} | {pr['P3']} | {pr['P4']} | {sum(1 for r in dic_rows if r['ACCION_BETTERSOFT'] == 'NO_CALCULAR_EN_UPLUS')} |", "",
             "Los cambios de v0.9 a v0.9.1 están en `REPORTE_CIERRE_GATE09_UPLUS_SIES_2026.md` (sección F).", "",
             "## Cambios v0.9.1 → v1.0", "", "| CAMPO | ATRIBUTO | V0.9.1 | V1.0 | CLASE | MOTIVO |", "|---|---|---|---|---|---|"]
    for c in changelog:
        lines.append(f"| {c['CAMPO']} | {c['ATRIBUTO']} | {c['V091']} | {c['V10']} | {c['CLASE']} | {c['MOTIVO']} |")
    lines += ["", f"**SIN_CAMBIO:** {len(sin_cambio)} filas del diccionario ({', '.join(sin_cambio)}).", "",
              "Clases usadas: SIN_CAMBIO, MODIFICADO, RECLASIFICADO, PENDIENTE_BETTERSOFT. No hay campos NUEVOS ni ELIMINADOS en v1.0."]
    (SUB / "CHANGELOG_COLUMNAS_BETTERSOFT_v0.9.1_A_v1.0.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    # pendiente institucional de rectificación (Fase 10.13)
    p3 = PENDIENTES[2]
    (SUB / "PENDIENTE_EVALUAR_RECTIFICACION_MU2026_C09_02_C09_03.md").write_text(f"""# PENDIENTE_EVALUAR_RECTIFICACION_MU2026_C09_02_C09_03

Estado: **PENDIENTE_DECISION_INSTITUCIONAL** · Registrado: {FECHA} · No forma parte del requerimiento Bettersoft.

En este gate no se rectificó nada, no se prepararon cargas correctivas y no se modificó Matrícula Unificada 2026.

## Universo afectado (MU 2026 pregrado cargada en PES, 08-05 y 11-05-2026)

| Resolución | Registros | Detalle |
|---|---|---|
| C09-02 Forma de ingreso 3 por cambio de jornada | {M['ct02']['pes_for3']} con FOR 3 | {f3.get('mismo_COD_CAR_distinta_JOR', 0)} con el mismo COD_CAR y otra jornada en MU 2025; {f3.get('mismo_COD_CAR_misma_JOR', 0)} con el mismo COD_CAR y la misma jornada; {f3.get('distinto_COD_CAR', 0)} con otro COD_CAR; {f3.get('sin_registro_MU2025', 0)} sin registro 2025. Los de mismo COD_CAR tenían FOR 1 en 2025. |
| C09-03 Forma de ingreso 2 por nombre "CONTINUIDAD" | {M['ct03']['pes_for2']} con FOR 2 | {f2.get('for_2025_11', 0)} con FOR 11 en 2025 en el mismo programa; {f2.get('for_2025_1', 0)} con FOR 1; {f2.get('sin_mismo_programa_MU2025', 0)} sin el mismo programa en 2025. |

## Naturaleza del problema

- C09-02: la regla oficial de Cambio Interno exige una carrera nueva (código de carrera distinto al del año anterior). Un cambio de jornada no lo prueba.
- C09-03: la forma 2 exige haber cursado un plan común o bachillerato. El nombre del programa no lo prueba, y la forma de ingreso del mismo estudiante al mismo programa cambió entre 2025 y 2026 sin un hecho que lo explique.

## Evidencia

- `REPORTE_CIERRE_GATE09_UPLUS_SIES_2026.md`, secciones C y D.
- `resultados/AUDITORIA_GATE09_CONTRADICCIONES_UPLUS_SIES_2026.xlsx`, hojas 01 y 02.
- `resultados/gate09/METRICAS_EVIDENCIA_GATE09.json` (claves `ct02` y `ct03`).
- Archivos cargados en PES, solo locales: `.local_restricted/requerimiento_bettersoft_uplus_2026/gate09_evidencia/matricula_unificada_2026/`.

## Impacto posible

Indicadores de cohorte, cambio interno y continuidad, y asignación o renovación de beneficios que dependen de la forma de ingreso y del año de origen. El impacto real depende de la regla de beneficios aplicable, que no se evaluó.

## Decisión institucional requerida

Definir si se solicita a SIES la corrección histórica de estos registros y con qué forma de ingreso correcta para cada uno. Esto requiere reconstruir los hechos de cada caso, no solo reemplazar un código. Según el Instructivo AC 2026, la corrección histórica se solicita formalmente a SIES y se hace efectiva en 2027.
""", encoding="utf-8")

    print("V1:", V10, "| changelog:", len(changelog), "| sin cambio:", len(sin_cambio), "| matriz:", len(matriz), pr)


if __name__ == "__main__":
    main()
