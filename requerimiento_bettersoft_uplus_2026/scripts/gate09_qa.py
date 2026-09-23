"""Gate 09: QA del diccionario v0.9.1, del registro de auditoría y del Word para el proveedor.

Escribe resultados/gate09/QA_GATE09.json. No modifica ningún entregable.
"""
import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.dom import minidom

import docx
import pandas as pd

SUB = Path(__file__).resolve().parents[1]
RES = SUB / "resultados"
V09 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026.xlsx"
V091 = RES / "DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v0.9.1_GATE09.xlsx"
AUD = RES / "AUDITORIA_GATE09_CONTRADICCIONES_UPLUS_SIES_2026.xlsx"
DOCX = RES / "REQUERIMIENTO_BETTERSOFT_REPORTE_REGULATORIO_UPLUS_v0.9.1_GATE09.docx"
ORIG = {  # SHA-256 de los entregables v0.9 registrados al iniciar el Gate 09
    "resultados/DICCIONARIO_MAESTRO_UPLUS_SIES_2026.xlsx": "94f6d9c7b0aa5e156360fe90e70a6cc6de5309d8a2e34beeb4e279770adb2885",
    "resultados/REQUERIMIENTO_BETTERSOFT_REPORTE_REGULATORIO_UPLUS.docx": "70a4c9f52312833206b0703bbf06de2faedab8cbd0060108cd9b51b5f9411482",
    "scripts/generar_diccionario_maestro_uplus_sies.py": "d156b93fdcdd9ca48f9e0e89b19c84fe6b5600be122f7dadb85fa53fa8e73705",
    "scripts/generar_requerimiento_bettersoft_docx.js": "68851d7128531dfc5fed95e2dd7a2c78fabe60a3f73a7c193af503900ae750ea",
    "README.md": "eaa25802bcf67e5a4eb6c64f6ea31268d4bf59402c8e19e95f67abd54cde6ef1",
}
ESTADOS_GATE = {"RESUELTA_REGLA_OFICIAL", "RESUELTA_DATO_REAL", "IMPLEMENTACION_HISTORICA_INCORRECTA", "NO_APLICA_AL_PROCESO",
                "NO_DETERMINABLE_DESDE_U_PLUS", "PENDIENTE_FALTA_FUENTE"}
PRIORIDADES = {"P1", "P2", "P3", "P4", "NO_APLICA_BETTERSOFT"}
CLASIF_CAT = {"OFICIAL", "OBSERVADO", "PENDIENTE", "NO_DETERMINABLE", "DEFAULT_TECNICO_NO_REGLA", "DESCARTADO_COMO_REGLA", "SIN_REVISION_GATE09"}
PROHIBIDOS = ["HIPOTESIS", "HIPÓTESIS", "hipótesis", "hipotesis", "C09-", "CT-0", "CT-1", "Gate", "GATE", ".py", ".tsv", ".json", "codigo_gobernanza",
              "decisión interna", "Decisión interna", "contradic", "Contradic", "PENDIENTE_DEMOSTRAR", "PUENTE", "scratchpad", "auditor", "Hoja1",
              "IMPLEMENTACION", "implementación histórica", "OBSERVADO", "DEFAULT_TECNICO", "SUPUESTO", "local_restricted"]
H1_ESPERADOS = ["1. Objetivo", "2. Alcance", "3. Principio del requerimiento", "4. Extractos solicitados", "5. Campos solicitados",
                "6. Maestro del código único SIES", "7. Datos históricos", "8. Catálogos de U+", "9. Datos que U+ no debe calcular",
                "10. Reglas de extracción", "11. Validaciones mínimas del reporte", "12. Consultas a Bettersoft", "Anexo. Cambios solicitados"]
RUT_RE = re.compile(r"(?<![\w.])\d{1,2}\.?\d{3}\.?\d{3}-[\dkK](?![\w])|(?<![\w])\d{7,8}(?![\w])")


PLACEHOLDERS = {"12345678", "123456789"}


def posibles_rut(texto):
    """Números con forma de RUN, excluyendo máscaras de perfil (solo 9), ejemplos ficticios y hashes de commit."""
    hallados = []
    for mm in RUT_RE.finditer(texto):
        dig = re.sub(r"\D", "", mm.group(0).split("-")[0])
        if set(dig) == {"9"} or dig in PLACEHOLDERS or texto[:mm.start()].rstrip().endswith("Commit"):
            continue
        hallados.append(mm.group(0))
    return hallados


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def celdas(path):
    for s, df in pd.read_excel(path, sheet_name=None, dtype=str).items():
        df = df.drop(columns=[c for c in df.columns if c == "bytes"])  # tamaños de archivo del manifiesto
        for v in df.fillna("").astype(str).values.ravel():
            yield s, v


def qa_xlsx():
    x = pd.ExcelFile(V091)
    hojas = {s: pd.read_excel(x, s, dtype=str).fillna("") for s in x.sheet_names}
    v09 = pd.ExcelFile(V09).sheet_names
    m, d, c, ct = hojas["04_MATRIZ_MAESTRA"], hojas["05_DICCIONARIO_BETTERSOFT"], hojas["06_CATALOGOS_HOMOLOGACIONES"], hojas["12_GATE09_RESOLUCION"]
    sol = d[d.ACCION_BETTERSOFT.isin(["SOLICITAR_HECHO", "SOLICITAR_MAESTRO"])]
    r = {
        "hojas": len(x.sheet_names), "hojas_v09_conservadas": all(s in x.sheet_names for s in v09), "hojas_nuevas": [s for s in x.sheet_names if s not in v09],
        "filas_por_hoja": {s: len(df) for s, df in hojas.items()},
        "matriz_filas": len(m), "matriz_id_duplicados": int(m.ID_FILA.duplicated().sum()),
        "matriz_campo_duplicado_por_subproceso": int(m.duplicated(["SUBPROCESO", "CAMPO_SIES"]).sum()),
        "matriz_filas_actualizadas_gate09": int((m.RESOLUCION_GATE09 != "").sum()),
        "diccionario_filas": len(d), "diccionario_id_duplicados": int(d.ID.duplicated().sum()),
        "procesos_matriz": sorted(m.PROCESO.unique().tolist()),
        "prioridades_invalidas": sorted(set(d.PRIORIDAD) - PRIORIDADES),
        "prioridades": d.PRIORIDAD.value_counts().to_dict(),
        "acciones": d.ACCION_BETTERSOFT.value_counts().to_dict(),
        "estados_gate09_invalidos": sorted(set(ct.ESTADO_GATE09) - ESTADOS_GATE),
        "contradicciones": len(ct), "contradicciones_por_estado": ct.ESTADO.value_counts().to_dict(),
        "contradicciones_por_estado_gate09": ct.ESTADO_GATE09.value_counts().to_dict(),
        "catalogo_clasificaciones_invalidas": sorted(set(c.CLASIFICACION_GATE09) - CLASIF_CAT),
        "catalogo_clasificaciones": c.CLASIFICACION_GATE09.value_counts().to_dict(),
        "matriz_filas_sin_fuente": int(((m.FUENTE_OFICIAL == "") & (m.FUENTE_INSTITUCIONAL == "") & (m.SCRIPT_O_EVIDENCIA == "")).sum()),
        "matriz_filas_sin_nivel_respaldo": int((m.NIVEL_RESPALDO == "").sum()),
        "diccionario_filas_sin_fuente": int(((d.EVIDENCIA == "") & (d.RESPALDO_NECESIDAD == "")).sum()),
        "diccionario_filas_sin_nivel_respaldo": int((d.NIVEL_RESPALDO_NECESIDAD == "").sum()),
        "P1_total": int((d.PRIORIDAD == "P1").sum()),
        "P1_NO_DEMOSTRADOS": int(((d.PRIORIDAD == "P1") & (d.P1_DEMOSTRADO != "SI")).sum()),
        "P1_disponibilidad_por_confirmar_o_nueva": int(((d.PRIORIDAD == "P1") & (d.DISPONIBILIDAD_UPLUS != "EXISTE")).sum()),
        "bettersoft_calculados_sin_regla": sol[sol.TIPO_OBTENCION.str.startswith(("B", "D"))].ID.tolist(),
        "supuestos_como_regla_oficial": int(((c.CLASIFICACION_GATE09 == "OFICIAL") & (c.LADO != "SIES")).sum()),
    }
    fugas = [(s, v[:40]) for p in (V091, AUD) for s, v in celdas(p) if posibles_rut(v)]
    r["posibles_datos_personales"] = fugas[:10]
    r["QA_XLSX"] = "APROBADO" if (r["hojas_v09_conservadas"] and not r["matriz_id_duplicados"] and not r["diccionario_id_duplicados"]
                                  and not r["prioridades_invalidas"] and not r["estados_gate09_invalidos"] and r["contradicciones"] == 13
                                  and not r["catalogo_clasificaciones_invalidas"] and not r["matriz_filas_sin_fuente"] and not r["diccionario_filas_sin_fuente"]
                                  and not r["diccionario_filas_sin_nivel_respaldo"] and r["P1_NO_DEMOSTRADOS"] == 0
                                  and not r["bettersoft_calculados_sin_regla"] and not r["supuestos_como_regla_oficial"] and not fugas) else "NO_APROBADO"
    return r


def qa_docx():
    r = {}
    with zipfile.ZipFile(DOCX) as z:
        r["zip_integro"] = z.testzip() is None
        partes = z.namelist()
        r["partes_requeridas"] = all(p in partes for p in ["[Content_Types].xml", "word/document.xml", "word/styles.xml", "word/_rels/document.xml.rels"])
        for p in [p for p in partes if p.endswith(".xml")]:
            minidom.parseString(z.read(p))
        r["xml_bien_formado"] = True
        xml = z.read("word/document.xml").decode("utf-8")
    r["tabla_contenido"] = "TOC" in xml
    r["orientacion_horizontal"] = 'w:orient="landscape"' in xml
    d = docx.Document(DOCX)
    h1 = [p.text for p in d.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    r["secciones_h1"] = h1
    r["secciones_esperadas_presentes"] = h1 == H1_ESPERADOS
    r["tablas"] = len(d.tables)
    # sección vacía: un H1 seguido directamente por otro H1 sin contenido (párrafo o tabla) entre ellos
    body = d.element.body
    vacias, actual, contenido = [], None, False
    for el in body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            estilo = el.find(".//{*}pStyle")
            texto = "".join(t.text or "" for t in el.iter("{*}t")).strip()
            if estilo is not None and estilo.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "Heading1":
                if actual and not contenido:
                    vacias.append(actual)
                actual, contenido = texto, False
            elif texto:
                contenido = True
        elif tag == "tbl":
            contenido = True
    if actual and not contenido:
        vacias.append(actual)
    r["secciones_vacias"] = vacias
    texto = "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])
    r["terminos_internos_encontrados"] = sorted({t for t in PROHIBIDOS if t in texto})
    r["posibles_datos_personales"] = posibles_rut(texto)[:10]
    req = json.loads((RES / "gate09/req_gate09.json").read_text(encoding="utf-8"))
    ids_sol = {x["ID"] for x in req["dic"]}
    ids_noc = {x["ID"] for x in req["no_calcular"]}
    r["contradiccion_interna_solicitar_y_no_calcular"] = sorted(ids_sol & ids_noc)
    p1 = sum(1 for x in req["dic"] if x["PRIORIDAD"] == "P1")
    r["conteo_p1_coherente"] = f"{p1} de prioridad P1" in texto
    r["QA_ESTRUCTURAL_DOCX"] = "APROBADO" if (r["zip_integro"] and r["partes_requeridas"] and r["secciones_esperadas_presentes"] and not vacias
                                              and not r["terminos_internos_encontrados"] and not r["posibles_datos_personales"]
                                              and not r["contradiccion_interna_solicitar_y_no_calcular"] and r["conteo_p1_coherente"]) else "NO_APROBADO"
    r["QA_VISUAL_DOCX"] = "PENDIENTE_REVISION_HUMANA"
    return r


def main():
    orig = {k: (sha(SUB / k) == v) for k, v in ORIG.items()}
    out = {"originales_v09_intactos": orig, "ORIGINALES_MODIFICADOS": "NO" if all(orig.values()) else "SI", "xlsx": qa_xlsx(), "docx": qa_docx()}
    (RES / "gate09/QA_GATE09.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(out, ensure_ascii=False, indent=1))
    return 0 if out["xlsx"]["QA_XLSX"] == "APROBADO" and out["docx"]["QA_ESTRUCTURAL_DOCX"] == "APROBADO" and out["ORIGINALES_MODIFICADOS"] == "NO" else 1


if __name__ == "__main__":
    sys.exit(main())
