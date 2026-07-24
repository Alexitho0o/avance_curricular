"""Entregable 2: Excel de respaldo con encabezados, tipado real y diccionario."""

import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent))  # raíz del monorepo (para 'common.*')
sys.path.insert(0, str(Path(__file__).parent.parent))       # procesos/ire_2026 (para 'src.*')

from datetime import datetime
from typing import List, Dict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from common.io_csv import leer_csv
from common.excel import configurar_hoja_base, guardar_libro
from src.schema import COLUMNAS, ESPECIFICACION, CATALOGO_TIPO_INFRAESTRUCTURA

RAIZ_PROCESO = Path(__file__).parent.parent
RUTA_OUTPUT = RAIZ_PROCESO / "output"
RUTA_HISTORICO = RAIZ_PROCESO / "data" / "historico" / "ire_historico.csv"

# Columnas donde el valor, si viene, es numérico real (no texto) en el Excel
COLUMNAS_NUMERICAS = {
    col for col, spec in ESPECIFICACION.items()
    if spec["tipo"] in ("numero", "numero_entero")
}


def _valor_tipado(columna: str, valor: str):
    """Convertir a número si la columna es numérica y el valor no está vacío."""
    if not valor or not valor.strip():
        return None
    if columna in COLUMNAS_NUMERICAS:
        try:
            numero = float(valor)
            return int(numero) if numero == int(numero) else numero
        except ValueError:
            return valor
    return valor


# Semáforo de estado (Mejora 2): colores y campos afectados en Miguel Claro 337.
# NOTA: el enunciado también sugiere un criterio genérico ("si la clave en
# parametros_2026.yaml es null -> amarillo, si tiene valor explícito -> verde"),
# pero ese criterio por sí solo pintaría de amarillo TOTAL_M2_TERRENO/EDIFICADOS/etc.
# de Miguel Claro (inmueble_permanente sigue en null en el YAML), contradiciendo
# la lista explícita de campos VERDE dada más abajo. Se usa la lista explícita
# como fuente de verdad: esos valores fueron confirmados directamente en el
# input (Miguel Claro es un domicilio nuevo, no hay "2025" del cual heredar),
# mientras que biblioteca/digital sí son copias literales de los valores 2025
# de O'Higgins. Ver informe final para el detalle de esta decisión.
ROJO = "FF0000"
AMARILLO = "FFFF00"
VERDE = "00B050"
GRIS = "D3D3D3"

MC_TIPO1_CAMPOS_VERDES = [
    "TOTAL_M2_TERRENO", "TOTAL_M2_EDIFICADOS", "TOTAL_SALAS_CLASES",
    "CAPACIDAD_SALAS_CLASES", "TOTAL_M2_SALAS_CLASES", "TOTAL_AUDITORIOS",
    "CAPACIDAD_AUDITORIOS", "TOTAL_M2_AUDITORIOS", "TOTAL_LABORATORIOS",
    "TOTAL_M2_LABORATORIOS", "TOTAL_TALLERES", "TOTAL_M2_TALLERES",
    "TOTAL_PC_NB_DISPONIBLE", "TOTAL_M2_CASINOS_CAFETERIAS", "TOTAL_M2_AREAS_VERDES",
]
MC_TIPO3_CAMPOS_AMARILLOS = [
    "TOTAL_M2_BIBLIOTECA", "TOTAL_M2_SALAS_LECTURA", "TOTAL_PROFESIONALES_BIBLIOTECA",
    "HORAS_PERSONAL_BIBLIOTECA", "TOTAL_TITULOS_DISPONIBLES",
    "TOTAL_VOLUMENES_DISPONIBLES", "TOTAL_SUSCRIPCIONES_REVISTAS",
]
MC_TIPO4_CAMPOS_AMARILLOS = [
    "TOTAL_TITULOS_LIBROS_DIGITALES", "TOTAL_SUSCRIPCIONES_DIGITALES", "TOTAL_BASE_DATOS",
]


def _pintar(ws, fila_idx: int, columna: str, color: str, texto_blanco: bool) -> None:
    col_idx = COLUMNAS.index(columna) + 1
    celda = ws.cell(row=fila_idx, column=col_idx)
    celda.fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
    celda.font = Font(color="FFFFFF" if texto_blanco else "000000")


def _aplicar_semaforo(ws, filas: List[Dict]) -> None:
    """Aplicar colores de estado a la hoja CARGA_2026 (ver notas arriba)."""
    for fila_idx, fila in enumerate(filas, start=2):
        tipo = fila.get("TIPO_INFRAESTRUCTURA", "").strip()
        direccion = fila.get("DIRECCION_INMUEBLE", "").strip()
        vigencia = fila.get("VIGENCIA", "").strip()

        if vigencia == "0":
            # Fila en eliminación (O'Higgins): gris completo
            for col_idx in range(1, len(COLUMNAS) + 1):
                ws.cell(row=fila_idx, column=col_idx).fill = PatternFill(
                    start_color=GRIS, end_color=GRIS, fill_type="solid"
                )
            continue

        if tipo == "1" and direccion == "MIGUEL CLARO 337":
            _pintar(ws, fila_idx, "FECHA_TERMINO", ROJO, texto_blanco=True)
            for campo in MC_TIPO1_CAMPOS_VERDES:
                if ws.cell(row=fila_idx, column=COLUMNAS.index(campo) + 1).value not in (None, ""):
                    _pintar(ws, fila_idx, campo, VERDE, texto_blanco=True)

        elif tipo == "3" and direccion == "MIGUEL CLARO 337":
            for campo in MC_TIPO3_CAMPOS_AMARILLOS:
                if ws.cell(row=fila_idx, column=COLUMNAS.index(campo) + 1).value not in (None, ""):
                    _pintar(ws, fila_idx, campo, AMARILLO, texto_blanco=False)

        elif tipo == "4" and direccion == "MIGUEL CLARO 337":
            for campo in MC_TIPO4_CAMPOS_AMARILLOS:
                if ws.cell(row=fila_idx, column=COLUMNAS.index(campo) + 1).value not in (None, ""):
                    _pintar(ws, fila_idx, campo, AMARILLO, texto_blanco=False)


def construir_hoja_carga(wb: Workbook, filas: List[Dict]) -> None:
    """Hoja CARGA_2026: datos con encabezados, tipado real y semáforo de estado."""
    ws = wb.active
    ws.title = "CARGA_2026"

    configurar_hoja_base(ws, COLUMNAS)

    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, columna in enumerate(COLUMNAS, start=1):
            valor = _valor_tipado(columna, fila.get(columna, ""))
            ws.cell(row=fila_idx, column=col_idx, value=valor)

    _aplicar_semaforo(ws, filas)

    # Leyenda
    fila_leyenda = len(filas) + 3
    leyenda = [
        ("■ ROJO     = pendiente de confirmar ANTES DE SUBIR A PES", ROJO, True),
        ("■ AMARILLO = heredado de 2025, confirmar con la unidad correspondiente", AMARILLO, False),
        ("■ VERDE    = confirmado para 2026", VERDE, True),
        ("■ GRIS     = registro en eliminación (VIGENCIA=0, no se mostrará en PES)", GRIS, False),
    ]
    for offset, (texto, color, _texto_blanco) in enumerate(leyenda):
        celda = ws.cell(row=fila_leyenda + offset, column=1, value=texto)
        celda.font = Font(italic=True, size=9)


def construir_hoja_diccionario(wb: Workbook) -> None:
    """Hoja DICCIONARIO: las 54 columnas con definición, códigos y especificación."""
    ws = wb.create_sheet("DICCIONARIO")

    encabezados = ["#", "CAMPO", "TIPO", "APLICA A TIPO", "OBLIGATORIO", "CATÁLOGO / NOTA"]
    configurar_hoja_base(ws, encabezados, anchos=[5, 35, 16, 16, 14, 60])

    catalogos_nota = {
        "TIPO_INFRAESTRUCTURA": "1=Uso Permanente, 2=Uso Restringido, 3=Biblioteca, 4=Digital, 5=Predios, 6=Plataformas",
        "SITUACION_TENENCIA": "1=Propio, 2=Arrendado, 3=Comodato, 4=Usufructo, 5=Leasing, 6=Otro",
        "USO_EXCLUSIVO": "1=Exclusivo, 2=Compartido",
        "UR_SITUACION_TENENCIA": "1=Arrendado, 2=Otra",
        "VIGENCIA": "0=Eliminar registro cargado, 1=Mantener",
    }

    for i, columna in enumerate(COLUMNAS, start=1):
        spec = ESPECIFICACION[columna]
        aplica = ", ".join(str(t) for t in sorted(spec["aplica"]))
        obligatorio = "Sí" if spec["obligatorio"] else "Condicional"
        nota = catalogos_nota.get(columna, "")

        ws.cell(row=i + 1, column=1, value=i)
        ws.cell(row=i + 1, column=2, value=columna)
        ws.cell(row=i + 1, column=3, value=spec["tipo"])
        ws.cell(row=i + 1, column=4, value=aplica)
        ws.cell(row=i + 1, column=5, value=obligatorio)
        ws.cell(row=i + 1, column=6, value=nota)


def _cargar_datos_2026_historico() -> Dict:
    """
    Cargar las magnitudes 2026 para la hoja HISTORICO desde sus fuentes reales:
    matrícula desde config/parametros_2026.yaml, magnitudes físicas y de
    plataformas desde la fila vigente (VIGENCIA=1) del domicilio actual
    (Miguel Claro 337) en data/input/ire_2026.csv, razón social desde
    config/institucion.yaml.
    """
    from src.config_loader import cargar_parametros, cargar_institucion

    ruta_input = RAIZ_PROCESO / "data" / "input" / "ire_2026.csv"
    _, filas_input = leer_csv(str(ruta_input), delimitador=';', encoding='utf-8')

    fila_tipo1 = fila_tipo4 = fila_tipo6 = None
    for f in filas_input:
        if f.get("VIGENCIA", "").strip() != "1":
            continue
        if f.get("DIRECCION_INMUEBLE", "").strip() != "MIGUEL CLARO 337":
            continue
        tipo = f.get("TIPO_INFRAESTRUCTURA", "").strip()
        if tipo == "1":
            fila_tipo1 = f
        elif tipo == "4":
            fila_tipo4 = f
        elif tipo == "6":
            fila_tipo6 = f

    def num(fila, campo):
        if fila is None:
            return None
        v = fila.get(campo, "").strip()
        if not v:
            return None
        try:
            n = float(v)
            return int(n) if n == int(n) else n
        except ValueError:
            return None

    parametros = cargar_parametros()
    institucion = cargar_institucion()

    return {
        "matricula": parametros.get("matricula_jornada_principal_2026"),
        "TOTAL_M2_TERRENO": num(fila_tipo1, "TOTAL_M2_TERRENO"),
        "TOTAL_M2_EDIFICADOS": num(fila_tipo1, "TOTAL_M2_EDIFICADOS"),
        "TOTAL_M2_AREAS_VERDES": num(fila_tipo1, "TOTAL_M2_AREAS_VERDES"),
        "TOTAL_M2_TALLERES": num(fila_tipo1, "TOTAL_M2_TALLERES"),
        "TOTAL_AUDITORIOS": num(fila_tipo1, "TOTAL_AUDITORIOS"),
        "TOTAL_PC_NB_DISPONIBLE": num(fila_tipo1, "TOTAL_PC_NB_DISPONIBLE"),
        "TOTAL_TALLERES": num(fila_tipo1, "TOTAL_TALLERES"),
        "TOTAL_SALAS_CLASES": num(fila_tipo1, "TOTAL_SALAS_CLASES"),
        "TOTAL_LABORATORIOS": num(fila_tipo1, "TOTAL_LABORATORIOS"),
        "TOTAL_TITULOS_LIBROS_DIGITALES": num(fila_tipo4, "TOTAL_TITULOS_LIBROS_DIGITALES"),
        "lms": fila_tipo6.get("SISTEMA_GESTION_APRENDIZAJES") if fila_tipo6 else None,
        "videoconferencia": fila_tipo6.get("SISTEMA_VIDEO_CONFERENCIA") if fila_tipo6 else None,
        "razon_social": institucion.get("razon_social"),
    }


def construir_hoja_historico(wb: Workbook, filas_historico: List[Dict]) -> None:
    """Hoja HISTORICO: comparativo 2023-2026 (2023-2025 según kpi.py, 2026 desde input+config)."""
    ws = wb.create_sheet("HISTORICO")

    from src.kpi import SERIE_HISTORICA

    encabezados = ["CONCEPTO", "2023", "2024", "2025", "2026"]
    configurar_hoja_base(ws, encabezados, anchos=[35, 15, 15, 15, 15])

    datos_2026 = _cargar_datos_2026_historico()

    conceptos = [
        ("Matrícula jornada principal", "matricula", "matricula"),
        ("M² de terreno", "m2_terreno", "TOTAL_M2_TERRENO"),
        ("M² construidos", "m2_edificados", "TOTAL_M2_EDIFICADOS"),
        ("M² áreas verdes y esparcimiento", "m2_areas_verdes", "TOTAL_M2_AREAS_VERDES"),
        ("M² de talleres y laboratorios", "m2_talleres", "TOTAL_M2_TALLERES"),
        ("Nº auditorios", "total_auditorios", "TOTAL_AUDITORIOS"),
        ("Nº computadores", "total_pc_nb_disponible", "TOTAL_PC_NB_DISPONIBLE"),
        ("Nº talleres", "total_talleres", "TOTAL_TALLERES"),
        ("Nº salas", "total_salas_clases", "TOTAL_SALAS_CLASES"),
        ("Laboratorios", "total_laboratorios", "TOTAL_LABORATORIOS"),
        ("Libros digitales (e-books)", "total_titulos_libros_digitales", "TOTAL_TITULOS_LIBROS_DIGITALES"),
        ("LMS", "lms", "lms"),
        ("Videoconferencia", "videoconferencia", "videoconferencia"),
        ("Razón social informada", "razon_social", "razon_social"),
    ]

    for i, (etiqueta, clave_historico, clave_2026) in enumerate(conceptos, start=2):
        ws.cell(row=i, column=1, value=etiqueta)
        for j, anio in enumerate([2023, 2024, 2025], start=2):
            valor = SERIE_HISTORICA[anio].get(clave_historico, "—")
            ws.cell(row=i, column=j, value=valor)
        valor_2026 = datos_2026.get(clave_2026)
        ws.cell(row=i, column=5, value=valor_2026 if valor_2026 is not None else "—")


def generar_xlsx(filas: List[Dict], fecha_sufijo: str = None) -> Path:
    """Generar el Excel de respaldo completo."""
    if fecha_sufijo is None:
        fecha_sufijo = datetime.now().strftime('%Y%m%d')

    RUTA_OUTPUT.mkdir(parents=True, exist_ok=True)
    ruta_salida = RUTA_OUTPUT / f"IRE_2026_respaldo_{fecha_sufijo}.xlsx"

    wb = Workbook()
    construir_hoja_carga(wb, filas)
    construir_hoja_diccionario(wb)

    _, filas_historico = leer_csv(str(RUTA_HISTORICO), delimitador=';', encoding='utf-8')
    construir_hoja_historico(wb, filas_historico)

    guardar_libro(wb, str(ruta_salida))
    return ruta_salida


def main():
    from src.build_csv import construir_dataset, validar_dataset_completo, ValidacionFallida
    from src.config_loader import ConfiguracionIncompleta

    print("=" * 70)
    print("IRE 2026 — Generación del Excel de respaldo")
    print("=" * 70)

    try:
        filas = construir_dataset()
    except ConfiguracionIncompleta as e:
        print("\n✗ PIPELINE DETENIDO — configuración incompleta\n")
        print(str(e))
        sys.exit(1)

    es_valido, mensajes = validar_dataset_completo(filas)
    if not es_valido:
        print(f"\n✗ VALIDACIÓN FALLIDA — {len(mensajes)} error(es):\n")
        for m in mensajes:
            print(f"  - {m}")
        raise ValidacionFallida(f"{len(mensajes)} error(es) de validación")

    ruta = generar_xlsx(filas)
    print(f"\n✓ Generado: {ruta}")


if __name__ == '__main__':
    main()
