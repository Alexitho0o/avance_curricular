#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Analisis sectorizado independiente para ENCUESTA_CARACTERIZACION + AC.

No modifica el pipeline productivo ni los Excel originales. Copia el libro de
encuesta y agrega hojas analiticas nuevas al final.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import traceback
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Literal, Sequence, cast

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from unidecode import unidecode

from caracterizacion_limesurvey import (
    LIMESURVEY_DEDUP_DISCARD_COLUMNS,
    LimeSurveyDeduplicationError,
    deduplicar_resultados_limesurvey,
    normalize_cross_value,
)


PY_PREFIX = "[SECTORIZADO_AC]"
PANEL_SHEET = "PANEL_SECTORIZADO_EJECUTIVO"
NULL_LABEL = "SIN INFORMACIÓN"
RESPONSE_STATES = ("SI", "INCOMPLETO", "NO")
B9_COLOR_HEADER = "1F4E78"
B9_COLOR_TITLE = "D9EAF7"
B9_COLOR_SUBTITLE = "EAF2F8"
B9_COLOR_KPI = "E2F0D9"
B9_COLOR_WARN = "FFF2CC"
B9_COLOR_GOOD = "C6E0B4"
B9_COLOR_BAD = "F4CCCC"
B9_COLOR_NEUTRAL = "F3F3F3"
B9_COLOR_BORDER = "D9D9D9"
EXPECTED_AC_COLS = [
    "_AC_ASI_INS_ANT",
    "_AC_ASI_APR_ANT",
    "_AC_ASI_INS_HIS",
    "_AC_ASI_APR_HIS",
    "_AC_PROM_PRI_SEM",
    "_AC_PROM_SEG_SEM",
    "_AC_NIV_ACA",
    "_AC_NIV_ACA_FUENTE",
]
AC_AVAILABILITY_COLS = [
    "_AC_ASI_INS_ANT",
    "_AC_ASI_APR_ANT",
    "_AC_ASI_INS_HIS",
    "_AC_ASI_APR_HIS",
    "_AC_PROM_PRI_SEM",
    "_AC_PROM_SEG_SEM",
    "_AC_NIV_ACA",
]
MANDATORY_AUDIT_SHEETS = [
    "AUDITORIA_UNIVERSO_SECTORIZADO",
    "AUDITORIA_DUPLICADOS_ENCUESTA",
    "AUDITORIA_DESCARTES_ENCUESTA",
    "AUDITORIA_CONFLICTOS_AC",
    "AUDITORIA_RECONCILIACION_PANEL",
]
NEW_SHEETS = [
    "SECTOR_RESUMEN_EJECUTIVO",
    "SECTOR_RESPUESTA",
    "SECTOR_CARRERA",
    "SECTOR_ANO_INGRESO",
    "SECTOR_NUEVO_ANTIGUO",
    "SECTOR_NIVEL_AC",
    "SECTOR_PROMEDIOS_AC",
    "SECTOR_PREGUNTAS_CLAVE",
    "BASE_SECTORIZADA",
    "AUDITORIA_SECTORES",
    "AUDITORIA_CRUCE_AC",
    *MANDATORY_AUDIT_SHEETS,
]
GRANULARITY_SHEETS = [
    "SECTOR_JORNADA",
    "SECTOR_JORNADA_NA",
    "CARRERA_JORNADA_NA",
    "ANO_INGRESO_JORNADA_NA",
    "NIVEL_AC_JORNADA_NA",
    "PROMEDIO_AC_JORNADA_NA",
    "MATRIZ_SECTOR_JORNADA_NA",
    "LECTURA_EJECUTIVA_SECTOR",
    "AUDITORIA_GRANULARIDAD",
]
ALL_OUTPUT_SHEETS = NEW_SHEETS + GRANULARITY_SHEETS
EXECUTIVE_SHEETS = {
    "SECTOR_RESUMEN_EJECUTIVO",
    "LECTURA_EJECUTIVA_SECTOR",
    "MATRIZ_SECTOR_JORNADA_NA",
}
AUDIT_SHEETS = {
    "AUDITORIA_SECTORES",
    "AUDITORIA_CRUCE_AC",
    "AUDITORIA_GRANULARIDAD",
    *MANDATORY_AUDIT_SHEETS,
}
DETAIL_SHEETS = {
    "BASE_SECTORIZADA",
    "SECTOR_PREGUNTAS_CLAVE",
}
SECTOR_CARRERAS_OFICIAL = {
    "Tecnología": [
        "Ingeniería en Ciberseguridad",
        "Ingeniería en Ciencia de Datos",
        "Ingeniería en Conectividad y Redes",
        "Ingeniería en Informática",
        "Ingeniería Industrial",
        "Ingeniería en Prevención de Riesgos",
        "Técnico en Ciencia de Datos",
        "Técnico en Ciberseguridad",
        "Técnico en Conectividad y Redes",
        "Técnico en Infraestructura Cloud",
        "Técnico en Programación y Análisis de Sistemas",
        "Técnico en Programación Computacional",
        "Técnico en Prevención de Riesgos",
    ],
    "Administración y Comercio": [
        "Administración Pública",
        "Contador Auditor",
        "Ingeniería en Administración de Empresas",
        "Ingeniería en Finanzas",
        "Ingeniería en Logística",
        "Ingeniería en Marketing Digital",
        "Ingeniería en Recursos Humanos",
        "Ingeniería en Seguridad Privada",
        "Contabilidad General",
        "Técnico en Administración de Empresas",
        "Técnico en Administración Pública",
        "Técnico en Finanzas",
        "Técnico en Logística",
        "Técnico en Marketing Digital",
        "Técnico en Recursos Humanos",
    ],
    "Salud": [
        "Naturopatía",
        "Técnico en Enfermería",
        "Técnico en Enfermería e Instrumentación Quirúrgica",
        "Técnico en Farmacia",
    ],
}

CARRERA_A_SECTOR_OFICIAL = {
    "INGENIERIA EN CIBERSEGURIDAD": "Tecnología",
    "INGENIERIA EN CIENCIA DE DATOS": "Tecnología",
    "INGENIERIA EN CONECTIVIDAD Y REDES": "Tecnología",
    "INGENIERIA EN INFORMATICA": "Tecnología",
    "INGENIERIA INDUSTRIAL": "Tecnología",
    "INGENIERIA EN PREVENCION DE RIESGOS": "Tecnología",
    "TECNICO EN CIENCIA DE DATOS": "Tecnología",
    "TECNICO EN CIBERSEGURIDAD": "Tecnología",
    "TECNICO EN CONECTIVIDAD Y REDES": "Tecnología",
    "TECNICO EN INFRAESTRUCTURA CLOUD": "Tecnología",
    "TECNICO EN PROGRAMACION Y ANALISIS DE SISTEMAS": "Tecnología",
    "TECNICO EN PROGRAMACION COMPUTACIONAL": "Tecnología",
    "TECNICO EN PREVENCION DE RIESGOS": "Tecnología",
    "ADMINISTRACION PUBLICA": "Administración y Comercio",
    "CONTADOR AUDITOR": "Administración y Comercio",
    "AUDITORIA": "Administración y Comercio",
    "INGENIERIA EN ADMINISTRACION DE EMPRESAS": "Administración y Comercio",
    "INGENIERIA EN FINANZAS": "Administración y Comercio",
    "INGENIERIA EN LOGISTICA": "Administración y Comercio",
    "INGENIERIA EN MARKETING DIGITAL": "Administración y Comercio",
    "INGENIERIA EN RECURSOS HUMANOS": "Administración y Comercio",
    "INGENIERIA EN SEGURIDAD PRIVADA": "Administración y Comercio",
    "CONTABILIDAD GENERAL": "Administración y Comercio",
    "TECNICO EN ADMINISTRACION DE EMPRESAS": "Administración y Comercio",
    "TECNICO EN ADMINISTRACION PUBLICA": "Administración y Comercio",
    "TECNICO EN FINANZAS": "Administración y Comercio",
    "TECNICO EN LOGISTICA": "Administración y Comercio",
    "TECNICO EN MARKETING DIGITAL": "Administración y Comercio",
    "TECNICO EN RECURSOS HUMANOS": "Administración y Comercio",
    "NATUROPATIA": "Salud",
    "TECNICO EN ENFERMERIA": "Salud",
    "TECNICO EN ENFERMERIA E INSTRUMENTACION QUIRURGICA": "Salud",
    "TECNICO EN FARMACIA": "Salud",
}

ALIAS_CARRERA_OFICIAL = {
    "INGENIERIA CIBERSEGURIDAD": "INGENIERIA EN CIBERSEGURIDAD",
    "INGENIERIA CIENCIA DE DATOS": "INGENIERIA EN CIENCIA DE DATOS",
    "INGENIERIA CONECTIVIDAD Y REDES": "INGENIERIA EN CONECTIVIDAD Y REDES",
    "INGENIERIA INFORMATICA": "INGENIERIA EN INFORMATICA",
    "INGENIERIA PREVENCION DE RIESGOS": "INGENIERIA EN PREVENCION DE RIESGOS",
    "PREVENCION DE RIESGOS": "INGENIERIA EN PREVENCION DE RIESGOS",
    "INGENIERIA ADMINISTRACION DE EMPRESAS": "INGENIERIA EN ADMINISTRACION DE EMPRESAS",
    "INGENIERIA FINANZAS": "INGENIERIA EN FINANZAS",
    "INGENIERIA LOGISTICA": "INGENIERIA EN LOGISTICA",
    "INGENIERIA MARKETING DIGITAL": "INGENIERIA EN MARKETING DIGITAL",
    "INGENIERIA RECURSOS HUMANOS": "INGENIERIA EN RECURSOS HUMANOS",
    "INGENIERIA SEGURIDAD PRIVADA": "INGENIERIA EN SEGURIDAD PRIVADA",
    "SEGURIDAD PRIVADA": "INGENIERIA EN SEGURIDAD PRIVADA",
    "ENFERMERIA": "TECNICO EN ENFERMERIA",
    "ENFERMERIA E INSTRUMENTACION QUIRURGICA": "TECNICO EN ENFERMERIA E INSTRUMENTACION QUIRURGICA",
    "FARMACIA": "TECNICO EN FARMACIA",
}

REGLA_CONTINUIDAD_OFICIAL = {
    "descripcion": (
        "Si una carrera contiene la palabra CONTINUIDAD, se elimina esa marca para "
        "clasificar por la carrera base. El sector lo define el nombre base visible "
        "en la imagen oficial."
    ),
    "ejemplos": {
        "CONTINUIDAD INGENIERIA INDUSTRIAL": "Tecnología",
        "CONTINUIDAD INGENIERIA CIENCIA DE DATOS": "Tecnología",
        "CONTINUIDAD EN PREVENCION DE RIESGOS": "Tecnología",
        "CONTINUIDAD INGENIERIA EN SEGURIDAD PRIVADA": "Administración y Comercio",
        "CONTINUIDAD INGENIERIA LOGISTICA": "Administración y Comercio",
        "CONTINUIDAD AUDITORIA": "Administración y Comercio",
        "CONTINUIDAD NATUROPATIA": "Salud",
    },
}


def log(msg: str) -> None:
    print(f"{PY_PREFIX} {msg}", flush=True)


def norm_txt(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", str(value))
    text = unidecode(text).upper().strip()
    text = re.sub(r"\s+", " ", text)
    return text


def norm_header(value: Any) -> str:
    """Normaliza encabezados sin asumir idioma, acentos ni puntuación."""
    return re.sub(r"[^A-Z0-9]+", "_", norm_txt(value)).strip("_")


class SectorizadoValidationError(RuntimeError):
    """Fallo bloqueante limitado a la generación del libro sectorizado."""

    def __init__(self, message: str, audit: pd.DataFrame | None = None):
        super().__init__(message)
        self.audit = audit.copy() if isinstance(audit, pd.DataFrame) else pd.DataFrame()


@dataclass
class DedupResult:
    resultados_unicos: pd.DataFrame
    auditoria: pd.DataFrame
    metricas: dict[str, int]
    token_col: str
    submitdate_col: str | None
    datestamp_col: str | None
    response_id_col: str | None
    descartes: pd.DataFrame = field(default_factory=pd.DataFrame)


@dataclass
class SectorizadoResult:
    base: pd.DataFrame
    auditoria_universo: pd.DataFrame
    auditoria_duplicados: pd.DataFrame
    auditoria_descartes_encuesta: pd.DataFrame
    auditoria_conflictos_ac: pd.DataFrame
    auditoria_reconciliacion_panel: pd.DataFrame
    metadata: dict[str, Any] = field(default_factory=dict)


AUDIT_UNIVERSE_COLUMNS = [
    "ejecución",
    "año",
    "período",
    "etapa",
    "fuente",
    "filas_antes",
    "filas_despues",
    "diferencia",
    "llave",
    "llaves_unicas",
    "llaves_nulas",
    "llaves_duplicadas",
    "cardinalidad_esperada",
    "cardinalidad_observada",
    "resultado",
    "observación",
]

AUDIT_DUPLICATE_COLUMNS = [
    "token",
    "cantidad_registros",
    "registros_completos",
    "registros_incompletos",
    "ID seleccionado",
    "fecha seleccionada",
    "criterio de selección",
    "IDs descartados",
    "estado final",
    "observación",
]

AUDIT_AC_COLUMNS = [
    "llave",
    "tipo de llave",
    "cantidad de candidatos",
    "años involucrados",
    "períodos involucrados",
    "carreras involucradas",
    "valores distintos",
    "criterio aplicado",
    "resultado",
    "requiere revisión",
]


def _as_key_list(key: str | Sequence[str]) -> list[str]:
    return [key] if isinstance(key, str) else list(key)


def _key_series(df: pd.DataFrame, key: str | Sequence[str]) -> pd.Series:
    keys = _as_key_list(key)
    if not keys or any(col not in df.columns for col in keys):
        return pd.Series("", index=df.index, dtype="object")
    if df.empty:
        return pd.Series("", index=df.index, dtype="object")
    frame = pd.DataFrame(index=df.index)
    for col in keys:
        frame[col] = df[col].map(lambda value: "" if pd.isna(value) else str(value).strip())
    valid = frame.apply(lambda row: all(value != "" for value in row), axis=1)
    joined = frame.astype(str).agg("\x1f".join, axis=1)
    return joined.where(valid, "")


def _key_stats(df: pd.DataFrame, key: str | Sequence[str]) -> tuple[int, int, int]:
    values = _key_series(df, key)
    valid = values[values.ne("")]
    counts = valid.value_counts(dropna=False)
    return int(valid.nunique()), int(values.eq("").sum()), int(counts.gt(1).sum())


def _observed_cardinality(left: pd.DataFrame, right: pd.DataFrame, key: str | Sequence[str]) -> str:
    left_values = _key_series(left, key)
    right_values = _key_series(right, key)
    left_dup = bool(left_values[left_values.ne("")].duplicated(keep=False).any())
    right_dup = bool(right_values[right_values.ne("")].duplicated(keep=False).any())
    if left_dup and right_dup:
        return "many_to_many"
    if left_dup:
        return "many_to_one"
    if right_dup:
        return "one_to_many"
    return "one_to_one"


class CardinalityAudit:
    """Bitácora de conservación de filas para todo el pipeline sectorizado."""

    def __init__(self, execution_id: str, anio: int, periodo: int):
        self.execution_id = str(execution_id)
        self.anio = int(anio)
        self.periodo = int(periodo)
        self.rows: list[dict[str, Any]] = []

    def add(
        self,
        *,
        etapa: str,
        fuente: str,
        filas_antes: int,
        filas_despues: int,
        llave: str | Sequence[str] | None,
        key_df: pd.DataFrame | None = None,
        cardinalidad_esperada: str = "sin_cambio_de_universo",
        cardinalidad_observada: str = "sin_cambio_de_universo",
        resultado: str = "OK",
        observacion: str = "",
    ) -> None:
        key_label = "" if llave is None else "+".join(_as_key_list(llave))
        if key_df is not None and llave is not None:
            unique, nulls, duplicates = _key_stats(key_df, llave)
        else:
            unique = nulls = duplicates = 0
        self.rows.append(
            {
                "ejecución": self.execution_id,
                "año": self.anio,
                "período": self.periodo,
                "etapa": etapa,
                "fuente": fuente,
                "filas_antes": int(filas_antes),
                "filas_despues": int(filas_despues),
                "diferencia": int(filas_despues - filas_antes),
                "llave": key_label,
                "llaves_unicas": unique,
                "llaves_nulas": nulls,
                "llaves_duplicadas": duplicates,
                "cardinalidad_esperada": cardinalidad_esperada,
                "cardinalidad_observada": cardinalidad_observada,
                "resultado": resultado,
                "observación": observacion,
            }
        )

    def frame(self) -> pd.DataFrame:
        return pd.DataFrame(self.rows, columns=AUDIT_UNIVERSE_COLUMNS)


def merge_preservando_filas(
    left: pd.DataFrame,
    right: pd.DataFrame,
    *,
    on: str | Sequence[str],
    audit: CardinalityAudit,
    etapa: str,
    fuente: str,
    validate: Literal["one_to_one", "many_to_one"] = "many_to_one",
    how: Literal["left"] = "left",
    suffixes: tuple[str, str] = ("", "_AUX"),
) -> pd.DataFrame:
    """Merge gobernado: nunca admite many-to-many ni cambio de filas."""
    if validate not in {"one_to_one", "many_to_one"}:
        raise ValueError("Solo se admiten validaciones one_to_one o many_to_one.")
    keys = _as_key_list(on)
    before = len(left)
    observed = _observed_cardinality(left, right, keys)
    right_null_keys = int(_key_series(right, keys).eq("").sum())
    if right_null_keys:
        observed = "llave_auxiliar_nula"
    invalid = right_null_keys > 0 or observed == "many_to_many" or (
        validate == "many_to_one" and observed == "one_to_many"
    ) or (validate == "one_to_one" and observed != "one_to_one")
    if invalid:
        audit.add(
            etapa=etapa,
            fuente=fuente,
            filas_antes=before,
            filas_despues=before,
            llave=keys,
            key_df=left,
            cardinalidad_esperada=validate,
            cardinalidad_observada=observed,
            resultado="ERROR",
            observacion=(
                "Cruce bloqueado antes del merge por cardinalidad no permitida; "
                f"llaves nulas en fuente auxiliar={right_null_keys}."
            ),
        )
        raise SectorizadoValidationError(
            f"[{etapa}] cardinalidad no permitida: esperada={validate}, observada={observed}",
            audit.frame(),
        )
    try:
        merged = left.merge(
            right,
            on=keys,
            how=how,
            validate=validate,
            suffixes=suffixes,
            sort=False,
        )
    except Exception as exc:
        audit.add(
            etapa=etapa,
            fuente=fuente,
            filas_antes=before,
            filas_despues=before,
            llave=keys,
            key_df=left,
            cardinalidad_esperada=validate,
            cardinalidad_observada=observed,
            resultado="ERROR",
            observacion=f"Merge rechazado: {exc}",
        )
        raise SectorizadoValidationError(f"[{etapa}] merge rechazado: {exc}", audit.frame()) from exc
    after = len(merged)
    result = "OK" if after == before else "ERROR"
    audit.add(
        etapa=etapa,
        fuente=fuente,
        filas_antes=before,
        filas_despues=after,
        llave=keys,
        key_df=merged,
        cardinalidad_esperada=validate,
        cardinalidad_observada=observed,
        resultado=result,
        observacion="Merge con validate explícito y conservación de filas.",
    )
    if after != before:
        raise SectorizadoValidationError(
            f"[{etapa}] el enriquecimiento cambió filas: {before} -> {after}",
            audit.frame(),
        )
    return merged


def _find_column(df: pd.DataFrame, explicit: str | None, aliases: Iterable[str]) -> str | None:
    if explicit:
        if explicit not in df.columns:
            raise SectorizadoValidationError(f"No existe la columna explícita {explicit!r}.")
        return explicit
    by_norm: dict[str, str] = {}
    for col in df.columns:
        by_norm.setdefault(norm_header(col), str(col))
    for alias in aliases:
        found = by_norm.get(norm_header(alias))
        if found is not None:
            return found
    alias_norms = [norm_header(alias) for alias in aliases]
    for col in df.columns:
        normalized = norm_header(col)
        if any(alias and alias in normalized for alias in alias_norms):
            return str(col)
    return None


def _clean_identifier(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", str(value)).strip()
    if not text or text.lower() in {"nan", "none", "null", "nat"}:
        return ""
    if re.fullmatch(r"-?\d+\.0", text):
        text = text[:-2]
    return text


def _normalizar_identificador_scalar(value: Any, key_kind: str, campo: str) -> str:
    if isinstance(value, (pd.Series, pd.Index, np.ndarray, list, tuple, dict, set)):
        raise SectorizadoValidationError(
            f"{campo} debe ser escalar; se recibio {type(value).__name__}."
        )
    return normalize_cross_value(value, key_kind)


def normalizar_identificadores_controlado(
    values: Any,
    *,
    key_kind: str = "token",
    campo: str = "identificador",
) -> Any:
    """Normaliza identificadores preservando texto y rechazando celdas no escalares."""
    if isinstance(values, pd.Series):
        return values.map(lambda value: _normalizar_identificador_scalar(value, key_kind, campo))
    if isinstance(values, pd.Index):
        return pd.Index(
            [_normalizar_identificador_scalar(value, key_kind, campo) for value in values],
            name=values.name,
            dtype="object",
        )
    if isinstance(values, np.ndarray):
        if values.ndim != 1:
            raise SectorizadoValidationError(
                f"{campo} debe ser un arreglo unidimensional; ndim={values.ndim}."
            )
        return np.array(
            [_normalizar_identificador_scalar(value, key_kind, campo) for value in values.tolist()],
            dtype=object,
        )
    if isinstance(values, list):
        return [_normalizar_identificador_scalar(value, key_kind, campo) for value in values]
    return _normalizar_identificador_scalar(values, key_kind, campo)


def normalizar_codcli_sectorizado(value: Any) -> str:
    return norm_txt(_clean_identifier(value)).replace(" ", "")


def normalizar_rut_con_dv(value: Any) -> str:
    text = _clean_identifier(value).upper().replace(".", "").replace(" ", "")
    return re.sub(r"[^0-9K]", "", text)


def _canonical_series(df: pd.DataFrame, aliases: Sequence[str], default: Any = "") -> pd.Series:
    col = _find_column(df, None, aliases)
    if col is None:
        return pd.Series(default, index=df.index)
    return df[col].copy()


def construir_universo_maestro(
    universo_vigente: pd.DataFrame,
    *,
    anio: int,
    periodo: int,
    fecha_ejecucion: str,
    audit: CardinalityAudit,
    key_kind: str = "token",
) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Copia el universo oficial; no lo filtra ni lo reconstruye."""
    if not isinstance(universo_vigente, pd.DataFrame):
        raise SectorizadoValidationError("El universo vigente oficial no es un DataFrame.", audit.frame())
    base = universo_vigente.copy(deep=True).reset_index(drop=True)
    initial_rows = len(base)
    codcli_raw = _canonical_series(base, ["CODCLI", "CODCLI_NORMALIZADO", "COD_CLI"])
    rut_raw = _canonical_series(
        base,
        ["RUT_NORMALIZADO_CON_DV", "RUT_NORMALIZADO", "RUT", "N_DOC", "DOCUMENTO"],
    )
    token_col = _find_column(
        base,
        None,
        ["TOKEN_NORMALIZADO", "_token_match", "_token_norm", "TOKEN", "CODIGO_ACCESO"],
    )
    if token_col is not None:
        token = normalizar_identificadores_controlado(
            base[token_col], key_kind=key_kind, campo="TOKEN universo"
        )
        token_source = pd.Series("TOKEN", index=base.index)
    else:
        token = pd.Series("", index=base.index, dtype="object")
        token_source = pd.Series("", index=base.index, dtype="object")
    rut_norm = rut_raw.map(normalizar_rut_con_dv)
    rut_token = normalizar_identificadores_controlado(
        rut_raw, key_kind="document", campo="RUT fallback universo"
    )
    missing_token = token.eq("")
    token = token.where(~missing_token, rut_token)
    token_source = token_source.where(~missing_token, "RUT_FALLBACK")

    base["CODCLI"] = codcli_raw.map(normalizar_codcli_sectorizado)
    base["TOKEN_NORMALIZADO"] = token
    base["RUT_NORMALIZADO"] = rut_norm
    base["FUENTE_TOKEN_NORMALIZADO"] = token_source
    canonical = {
        "CARRERA": ["NOMBRE_L", "CARRERA", "NOMBRE_CARRERA"],
        "CODCARPR": ["CODCARPR", "CODCARR", "COD_CARRERA"],
        "JORNADA": ["JORNADA"],
        "SEDE": ["SEDE", "NOMBRE_SEDE"],
        "ANOINGRESO": ["ANOINGRESO", "ANIO_INGRESO", "AÑO_INGRESO"],
        "PERIODOINGRESO": ["PERIODOINGRESO", "PERIODO_INGRESO"],
        "NIVEL": ["NIVEL", "NIVEL_ACADEMICO"],
        "ESTADO_VIGENCIA": ["ESTADOACADEMICO", "ESTADO_VIGENCIA", "ESTADO"],
    }
    for target, aliases in canonical.items():
        source = _canonical_series(base, aliases, np.nan)
        base[target] = source
    base["CARRERA_NORMALIZADA"] = base["CARRERA"].map(norm_txt)
    base["ANIO_ANALISIS"] = int(anio)
    base["PERIODO_ANALISIS"] = int(periodo)
    base["FECHA_EJECUCION"] = str(fecha_ejecucion)

    for aliases, expected, label in [
        (["ANOMATRICULA", "ANIO", "AÑO"], int(anio), "año"),
        (["PERIODOMATRICULA", "PERIODO"], int(periodo), "período"),
    ]:
        source_col = _find_column(universo_vigente, None, aliases)
        if source_col is None:
            continue
        values = pd.to_numeric(universo_vigente[source_col], errors="coerce").dropna().astype(int)
        mismatches = int(values.ne(expected).sum())
        if mismatches:
            audit.add(
                etapa="validar_parametros_universo",
                fuente="DF_VIGENTES",
                filas_antes=initial_rows,
                filas_despues=initial_rows,
                llave="CODCLI",
                key_df=base,
                resultado="ERROR",
                observacion=f"{mismatches} filas no corresponden al {label} parametrizado {expected}.",
            )
            raise SectorizadoValidationError(
                f"El universo oficial contiene {mismatches} filas fuera del {label} {expected}.",
                audit.frame(),
            )

    technical_columns = [
        "RUT_NORMALIZADO",
        "CARRERA_NORMALIZADA",
        "JORNADA",
        "SEDE",
        "ANOINGRESO",
        "PERIODOINGRESO",
    ]
    keys: list[str] = []
    key_types: list[str] = []
    for _, row in base.iterrows():
        if row["CODCLI"]:
            keys.append(f"CODCLI|{row['CODCLI']}")
            key_types.append("CODCLI")
        elif row["TOKEN_NORMALIZADO"]:
            keys.append(f"TOKEN|{row['TOKEN_NORMALIZADO']}")
            key_types.append("TOKEN_FALLBACK")
        elif row["RUT_NORMALIZADO"]:
            keys.append(f"RUT|{row['RUT_NORMALIZADO']}")
            key_types.append("RUT_FALLBACK")
        else:
            payload = "|".join(_clean_identifier(row.get(col, "")) for col in technical_columns)
            digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
            keys.append(f"TECNICA|{digest}")
            key_types.append("LLAVE_TECNICA_COMPUESTA")
    base["LLAVE_ESTUDIANTE"] = keys
    base["TIPO_LLAVE_ESTUDIANTE"] = key_types
    duplicated = base["LLAVE_ESTUDIANTE"].duplicated(keep=False)
    if bool(duplicated.any()):
        audit.add(
            etapa="validar_unicidad_universo",
            fuente="DF_VIGENTES",
            filas_antes=initial_rows,
            filas_despues=initial_rows,
            llave="LLAVE_ESTUDIANTE",
            key_df=base,
            cardinalidad_esperada="one_to_one",
            cardinalidad_observada="one_to_many",
            resultado="ERROR",
            observacion=f"{int(duplicated.sum())} filas tienen llave principal repetida.",
        )
        raise SectorizadoValidationError(
            "La llave principal del universo vigente no es única.",
            audit.frame(),
        )
    audit.add(
        etapa="universo_maestro_oficial",
        fuente="DF_VIGENTES",
        filas_antes=initial_rows,
        filas_despues=len(base),
        llave="LLAVE_ESTUDIANTE",
        key_df=base,
        cardinalidad_esperada="one_to_one",
        cardinalidad_observada="one_to_one",
        observacion="Copia exacta del universo oficial; no se aplicaron filtros ni deduplicación.",
    )
    metadata = {
        "total_vigentes": initial_rows,
        "llave_principal": "LLAVE_ESTUDIANTE",
        "casos_fallback_llave_universo": int(pd.Series(key_types).ne("CODCLI").sum()),
        "casos_token_desde_rut": int(token_source.eq("RUT_FALLBACK").sum()),
    }
    return base, metadata


def deduplicar_resultados_encuesta(
    resultados: pd.DataFrame,
    *,
    token_col: str | None = None,
    submitdate_col: str | None = None,
    datestamp_col: str | None = None,
    response_id_col: str | None = None,
    key_kind: str = "token",
) -> DedupResult:
    """Adapta el deduplicador institucional sin cambiar el contrato sectorizado."""
    try:
        canonical = deduplicar_resultados_limesurvey(
            resultados,
            token_col=token_col,
            submitdate_col=submitdate_col,
            datestamp_col=datestamp_col,
            response_id_col=response_id_col,
            key_kind=key_kind,
        )
    except LimeSurveyDeduplicationError as exc:
        raise SectorizadoValidationError(str(exc)) from exc

    return DedupResult(
        resultados_unicos=canonical.resultados_unicos,
        auditoria=canonical.auditoria.reindex(columns=AUDIT_DUPLICATE_COLUMNS),
        metricas=dict(canonical.metricas),
        token_col=canonical.token_col,
        submitdate_col=canonical.submitdate_col,
        datestamp_col=canonical.datestamp_col,
        response_id_col=canonical.response_id_col,
        descartes=canonical.descartes.reindex(
            columns=LIMESURVEY_DEDUP_DISCARD_COLUMNS
        ),
    )


def validar_resultados_canonicos_precomputados(
    resultados_unicos: pd.DataFrame,
    *,
    metricas: dict[str, int],
    auditoria: pd.DataFrame,
    descartes: pd.DataFrame | None = None,
    token_col: str,
    submitdate_col: str | None = None,
    datestamp_col: str | None = None,
    response_id_col: str | None = None,
    key_kind: str = "token",
) -> DedupResult:
    """Consume el resultado institucional ya deduplicado sin deduplicarlo otra vez.

    El frame canónico completo se usa para cruces; las métricas y auditorías
    conservan la procedencia de la fuente física original. Esto evita que un
    consumidor downstream convierta artificialmente ``filas brutas`` en
    ``filas únicas`` o pierda los tokens no asociados al maestro vigente.
    """
    if not isinstance(resultados_unicos, pd.DataFrame):
        raise SectorizadoValidationError(
            "Los resultados canónicos precomputados no son un DataFrame."
        )
    requeridas = {
        "_TOKEN_RESULTADO",
        "ESTADO_RESPUESTA_RESULTADO",
    }
    faltantes = sorted(requeridas - set(resultados_unicos.columns))
    if faltantes:
        raise SectorizadoValidationError(
            "Resultados canónicos sin columnas requeridas: " + ", ".join(faltantes)
        )
    canonical = resultados_unicos.copy(deep=True)
    # Excel puede inferir un token compuesto solo por dígitos como entero.
    # Restituimos el mismo dominio textual institucional antes de cualquier
    # merge; esto no vuelve a seleccionar filas ni altera la deduplicación.
    canonical["_TOKEN_RESULTADO"] = normalizar_identificadores_controlado(
        canonical["_TOKEN_RESULTADO"],
        key_kind=key_kind,
        campo="_TOKEN_RESULTADO",
    )
    tokens = canonical["_TOKEN_RESULTADO"].fillna("").astype(str).str.strip()
    if tokens.eq("").any() or not tokens.is_unique:
        raise SectorizadoValidationError(
            "Los resultados canónicos precomputados no son únicos y no nulos por token."
        )
    if not canonical["ESTADO_RESPUESTA_RESULTADO"].isin({"SI", "INCOMPLETO"}).all():
        raise SectorizadoValidationError(
            "Los resultados canónicos contienen estados fuera de SI/INCOMPLETO."
        )

    metricas_norm = {str(k): int(v) for k, v in dict(metricas).items()}
    metricas_norm.setdefault("filas_resultados_unicos", len(canonical))
    metricas_norm.setdefault("tokens_resultados_unicos", int(tokens.nunique()))
    if metricas_norm["filas_resultados_unicos"] != len(canonical):
        raise SectorizadoValidationError(
            "La métrica filas_resultados_unicos no coincide con el frame canónico."
        )
    if metricas_norm["tokens_resultados_unicos"] != int(tokens.nunique()):
        raise SectorizadoValidationError(
            "La métrica tokens_resultados_unicos no coincide con el frame canónico."
        )
    filas_brutas = metricas_norm.get("filas_resultados_brutos", len(canonical))
    if filas_brutas < len(canonical):
        raise SectorizadoValidationError(
            "La procedencia declara menos filas brutas que resultados canónicos."
        )
    descartes_frame = (
        descartes.copy(deep=True)
        if isinstance(descartes, pd.DataFrame)
        else pd.DataFrame(columns=LIMESURVEY_DEDUP_DISCARD_COLUMNS)
    )
    filas_descartadas = metricas_norm.get(
        "filas_descartadas_totales", len(descartes_frame)
    )
    if filas_brutas != len(canonical) + filas_descartadas:
        raise SectorizadoValidationError(
            "La procedencia no reconcilia: filas brutas != canónicas + descartadas."
        )
    if len(descartes_frame) != filas_descartadas:
        raise SectorizadoValidationError(
            "La auditoría de descartes no coincide con filas_descartadas_totales."
        )
    audit_frame = auditoria.reindex(columns=AUDIT_DUPLICATE_COLUMNS).copy()
    if audit_frame.empty and len(canonical):
        raise SectorizadoValidationError(
            "Falta auditoría de selección para resultados canónicos precomputados."
        )
    return DedupResult(
        resultados_unicos=canonical,
        auditoria=audit_frame,
        metricas=metricas_norm,
        token_col=token_col,
        submitdate_col=submitdate_col,
        datestamp_col=datestamp_col,
        response_id_col=response_id_col,
        descartes=descartes_frame.reindex(
            columns=LIMESURVEY_DEDUP_DISCARD_COLUMNS
        ),
    )


def _distinct_values(series: pd.Series) -> list[Any]:
    values: list[Any] = []
    seen: set[str] = set()
    for value in series.tolist():
        if pd.isna(value):
            continue
        clean = _clean_identifier(value)
        if clean == "":
            continue
        marker = f"{type(value).__name__}:{clean}"
        if marker not in seen:
            values.append(value)
            seen.add(marker)
    return values


def _join_distinct(series: pd.Series) -> str:
    values = sorted({_clean_identifier(value) for value in series.tolist() if _clean_identifier(value)})
    return " | ".join(values)


def _consolidar_grupo_ac(
    group: pd.DataFrame,
    *,
    key_value: str,
    key_type: str,
    ac_cols: Sequence[str],
    anio: int,
    periodo: int,
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    out: dict[str, Any] = {"_AC_KEY": key_value}
    current_mask = group["_AC_ANIO"].eq(int(anio)) & group["_AC_PERIODO"].eq(int(periodo))
    current = group[current_mask].copy()
    conflicts: dict[str, list[str]] = {}
    criteria: set[str] = set()
    unresolved = False
    for col in ac_cols:
        all_values = _distinct_values(group[col])
        current_values = _distinct_values(current[col]) if not current.empty else []
        if len(all_values) <= 1:
            out[col] = all_values[0] if all_values else np.nan
            criteria.add("VALORES_EQUIVALENTES")
            continue
        conflicts[col] = [_clean_identifier(value) for value in all_values]
        if len(current_values) == 1:
            out[col] = current_values[0]
            criteria.add("PRIORIZA_ANIO_PERIODO_ACTUAL")
        else:
            out[col] = np.nan
            unresolved = True
            criteria.add("CONFLICTO_NO_RESUELTO")
    audit_row: dict[str, Any] | None = None
    if len(group) > 1 or conflicts:
        if unresolved:
            result = "CONFLICTO_NO_RESUELTO"
        elif conflicts:
            result = "CONFLICTO_RESUELTO_POR_PERIODO_ACTUAL"
        else:
            result = "CONSOLIDADO_EQUIVALENTE"
        audit_row = {
            "llave": key_value,
            "tipo de llave": key_type,
            "cantidad de candidatos": int(len(group)),
            "años involucrados": _join_distinct(group["_AC_ANIO"]),
            "períodos involucrados": _join_distinct(group["_AC_PERIODO"]),
            "carreras involucradas": _join_distinct(group["_AC_CARRERA"]),
            "valores distintos": json.dumps(conflicts, ensure_ascii=False, sort_keys=True),
            "criterio aplicado": " | ".join(sorted(criteria)),
            "resultado": result,
            "requiere revisión": "SI" if unresolved else "NO",
        }
    return out, audit_row


def reducir_fuente_ac(
    fuente_ac: pd.DataFrame | None,
    *,
    universo: pd.DataFrame,
    anio: int,
    periodo: int,
    audit: CardinalityAudit,
) -> tuple[dict[str, pd.DataFrame], pd.DataFrame, dict[str, Any]]:
    """Reduce AC a llaves únicas antes de cualquier incorporación al universo."""
    empty_tables = {
        "codcli": pd.DataFrame(columns=["CODCLI", *EXPECTED_AC_COLS]),
        "token": pd.DataFrame(columns=["TOKEN_NORMALIZADO", "_AC_CARRERA_FUENTE", *EXPECTED_AC_COLS]),
        "rut_carrera": pd.DataFrame(
            columns=["RUT_NORMALIZADO", "CARRERA_NORMALIZADA", *EXPECTED_AC_COLS]
        ),
    }
    if fuente_ac is None or not isinstance(fuente_ac, pd.DataFrame) or fuente_ac.empty:
        audit.add(
            etapa="reducir_fuente_ac",
            fuente="SIN_FUENTE_AC",
            filas_antes=0,
            filas_despues=0,
            llave="CODCLI",
            key_df=empty_tables["codcli"],
            cardinalidad_esperada="many_to_one",
            cardinalidad_observada="sin_datos",
            observacion="No se recibió fuente AC; el universo se conserva sin antecedentes AC.",
        )
        return empty_tables, pd.DataFrame(columns=AUDIT_AC_COLUMNS), {
            "conflictos_ac_pendientes": 0,
            "columnas_ac_disponibles": [],
        }

    src = fuente_ac.copy(deep=True).reset_index(drop=True)
    original_rows = len(src)
    codcli_raw = _canonical_series(src, ["CODCLI", "CODCLI_NORMALIZADO", "COD_CLI"])
    rut_raw = _canonical_series(src, ["RUT_NORMALIZADO_CON_DV", "RUT_NORMALIZADO", "RUT", "N_DOC"])
    token_raw = _canonical_series(src, ["TOKEN_NORMALIZADO", "_token_match", "_token_norm", "TOKEN"])
    carrera_raw = _canonical_series(src, ["NOMBRE_L", "CARRERA", "NOMBRE_CARRERA"])
    src["_AC_CODCLI"] = codcli_raw.map(normalizar_codcli_sectorizado)
    src["_AC_RUT"] = rut_raw.map(normalizar_rut_con_dv)
    src["_AC_TOKEN"] = normalizar_identificadores_controlado(
        token_raw, key_kind="token", campo="TOKEN fuente AC"
    )
    src["_AC_TOKEN"] = src["_AC_TOKEN"].where(
        src["_AC_TOKEN"].ne(""),
        normalizar_identificadores_controlado(
            rut_raw, key_kind="document", campo="RUT fallback fuente AC"
        ),
    )
    src["_AC_CARRERA"] = carrera_raw.map(norm_txt)
    src["_AC_ANIO"] = pd.to_numeric(
        _canonical_series(src, ["ANIO", "AÑO", "ANOMATRICULA"], np.nan), errors="coerce"
    )
    src["_AC_PERIODO"] = pd.to_numeric(
        _canonical_series(src, ["PERIODO", "PERIODOMATRICULA"], np.nan), errors="coerce"
    )
    ac_cols = [col for col in EXPECTED_AC_COLS if col in src.columns]
    for col in EXPECTED_AC_COLS:
        if col not in src.columns:
            src[col] = np.nan
    for col in EXPECTED_AC_COLS[:-1]:
        src[col] = pd.to_numeric(src[col], errors="coerce")
    relevant = (
        src["_AC_CODCLI"].isin(set(universo["CODCLI"]) - {""})
        | src["_AC_TOKEN"].isin(set(universo["TOKEN_NORMALIZADO"]) - {""})
        | src["_AC_RUT"].isin(set(universo["RUT_NORMALIZADO"]) - {""})
    )
    candidates = src[relevant].copy()
    conflict_rows: list[dict[str, Any]] = []

    def reduce_by(source: pd.DataFrame, key_cols: list[str], key_type: str) -> pd.DataFrame:
        valid_mask = source[key_cols].apply(
            lambda col: col.map(lambda value: _clean_identifier(value) != "")
        ).all(axis=1)
        valid_source = source[valid_mask].copy()
        rows: list[dict[str, Any]] = []
        if valid_source.empty:
            return pd.DataFrame()
        grouper: str | list[str] = key_cols[0] if len(key_cols) == 1 else key_cols
        for key_values, group in valid_source.groupby(grouper, sort=False, dropna=False):
            values_tuple: tuple[Any, ...] = (
                (key_values,)
                if len(key_cols) == 1
                else tuple(cast(Sequence[Any], key_values))
            )
            key_value = "|".join(_clean_identifier(value) for value in values_tuple)
            row, conflict = _consolidar_grupo_ac(
                group,
                key_value=key_value,
                key_type=key_type,
                ac_cols=EXPECTED_AC_COLS,
                anio=anio,
                periodo=periodo,
            )
            for col, value in zip(key_cols, values_tuple):
                row[col] = value
            row["_AC_CARRERA_FUENTE"] = (
                _distinct_values(group["_AC_CARRERA"])[0]
                if len(_distinct_values(group["_AC_CARRERA"])) == 1
                else ""
            )
            row["_AC_CODCLI_CANDIDATOS"] = int(group["_AC_CODCLI"].replace("", np.nan).nunique())
            rows.append(row)
            if conflict is not None and (
                key_type == "CODCLI" or bool(group["_AC_CODCLI"].eq("").all())
            ):
                conflict_rows.append(conflict)
        return pd.DataFrame(rows)

    by_codcli = reduce_by(candidates, ["_AC_CODCLI"], "CODCLI")
    if by_codcli.empty:
        codcli = empty_tables["codcli"].copy()
    else:
        codcli = by_codcli.rename(columns={"_AC_CODCLI": "CODCLI"})
        codcli = codcli[["CODCLI", *EXPECTED_AC_COLS]].copy()
    by_token = reduce_by(candidates, ["_AC_TOKEN"], "TOKEN_NORMALIZADO")
    if by_token.empty:
        token = empty_tables["token"].copy()
    else:
        token = by_token[by_token["_AC_CODCLI_CANDIDATOS"].le(1)].copy()
        token = token.rename(columns={"_AC_TOKEN": "TOKEN_NORMALIZADO"})
        token = token[["TOKEN_NORMALIZADO", "_AC_CARRERA_FUENTE", *EXPECTED_AC_COLS]].copy()
    by_rut_carrera = reduce_by(candidates, ["_AC_RUT", "_AC_CARRERA"], "RUT+CARRERA")
    if by_rut_carrera.empty:
        rut_carrera = empty_tables["rut_carrera"].copy()
    else:
        rut_carrera = by_rut_carrera.rename(
            columns={"_AC_RUT": "RUT_NORMALIZADO", "_AC_CARRERA": "CARRERA_NORMALIZADA"}
        )
        rut_carrera = rut_carrera[["RUT_NORMALIZADO", "CARRERA_NORMALIZADA", *EXPECTED_AC_COLS]].copy()

    for table_name, table, key in [
        ("CODCLI", codcli, "CODCLI"),
        ("TOKEN", token, "TOKEN_NORMALIZADO"),
        ("RUT+CARRERA", rut_carrera, ["RUT_NORMALIZADO", "CARRERA_NORMALIZADA"]),
    ]:
        key_values = _key_series(table, key)
        if bool(key_values[key_values.ne("")].duplicated(keep=False).any()):
            raise SectorizadoValidationError(
                f"La reducción AC no quedó única para {table_name}.", audit.frame()
            )
    audit.add(
        etapa="reducir_fuente_ac",
        fuente="TRAYECTORIA_POR_PERIODO_AC",
        filas_antes=original_rows,
        filas_despues=len(codcli),
        llave="CODCLI",
        key_df=codcli,
        cardinalidad_esperada="many_to_one_reduction",
        cardinalidad_observada="many_to_one_reduction",
        observacion=(
            f"Fuente AC reducida antes del cruce; candidatos relevantes={len(candidates)}; "
            f"columnas AC presentes={', '.join(ac_cols) if ac_cols else 'ninguna'}."
        ),
    )
    conflicts_df = pd.DataFrame(conflict_rows, columns=AUDIT_AC_COLUMNS)
    pending = int(conflicts_df["requiere revisión"].eq("SI").sum()) if not conflicts_df.empty else 0
    return {"codcli": codcli, "token": token, "rut_carrera": rut_carrera}, conflicts_df, {
        "conflictos_ac_pendientes": pending,
        "columnas_ac_disponibles": ac_cols,
        "filas_ac_candidatas": int(len(candidates)),
    }


def incorporar_ac(
    base: pd.DataFrame,
    tables: dict[str, pd.DataFrame],
    *,
    audit: CardinalityAudit,
) -> tuple[pd.DataFrame, int]:
    """Incorpora AC por CODCLI y fallbacks no ambiguos sin cambiar el universo."""
    out = base.copy()
    for col in EXPECTED_AC_COLS:
        if col in out.columns:
            out.drop(columns=[col], inplace=True)
    primary = tables.get("codcli", pd.DataFrame()).copy()
    if not primary.empty:
        primary["_AC_MATCH_CODCLI"] = True
    else:
        primary = pd.DataFrame(columns=["CODCLI", *EXPECTED_AC_COLS, "_AC_MATCH_CODCLI"])
    out = merge_preservando_filas(
        out,
        primary,
        on="CODCLI",
        audit=audit,
        etapa="enriquecer_ac_codcli",
        fuente="AC_REDUCIDA_CODCLI",
        validate="many_to_one",
    )
    out["FUENTE_MATCH_AC"] = np.where(out["_AC_MATCH_CODCLI"].fillna(False), "CODCLI", "")

    unmatched = out["FUENTE_MATCH_AC"].eq("")
    token_counts = out.loc[unmatched & out["TOKEN_NORMALIZADO"].ne(""), "TOKEN_NORMALIZADO"].value_counts()
    out["_AC_TOKEN_FALLBACK_KEY"] = out["TOKEN_NORMALIZADO"].where(
        unmatched & out["TOKEN_NORMALIZADO"].map(token_counts).eq(1), ""
    )
    token_table = tables.get("token", pd.DataFrame()).copy()
    if not token_table.empty:
        token_table = token_table.rename(columns={"TOKEN_NORMALIZADO": "_AC_TOKEN_FALLBACK_KEY"})
        rename = {col: f"{col}__TOKEN" for col in EXPECTED_AC_COLS}
        token_table.rename(columns=rename, inplace=True)
        token_table["_AC_MATCH_TOKEN"] = True
    else:
        token_table = pd.DataFrame(
            columns=["_AC_TOKEN_FALLBACK_KEY", "_AC_CARRERA_FUENTE", "_AC_MATCH_TOKEN", *[f"{c}__TOKEN" for c in EXPECTED_AC_COLS]]
        )
    out = merge_preservando_filas(
        out,
        token_table,
        on="_AC_TOKEN_FALLBACK_KEY",
        audit=audit,
        etapa="enriquecer_ac_token_fallback",
        fuente="AC_REDUCIDA_TOKEN",
        validate="many_to_one",
    )
    carrera_compatible = (
        out.get("_AC_CARRERA_FUENTE", pd.Series("", index=out.index)).fillna("").eq("")
        | out["CARRERA_NORMALIZADA"].fillna("").eq("")
        | out.get("_AC_CARRERA_FUENTE", pd.Series("", index=out.index)).fillna("").eq(
            out["CARRERA_NORMALIZADA"].fillna("")
        )
    )
    apply_token = (
        out["FUENTE_MATCH_AC"].eq("")
        & out["_AC_MATCH_TOKEN"].astype("boolean").fillna(False)
        & carrera_compatible
    )
    for col in EXPECTED_AC_COLS:
        candidate = f"{col}__TOKEN"
        out[col] = out[col].where(~apply_token, out[candidate])
    out.loc[apply_token, "FUENTE_MATCH_AC"] = "TOKEN_FALLBACK"

    unmatched = out["FUENTE_MATCH_AC"].eq("")
    composite = out["RUT_NORMALIZADO"].fillna("") + "|" + out["CARRERA_NORMALIZADA"].fillna("")
    composite = composite.where(
        out["RUT_NORMALIZADO"].fillna("").ne("") & out["CARRERA_NORMALIZADA"].fillna("").ne(""), ""
    )
    composite_counts = composite[unmatched & composite.ne("")].value_counts()
    out["_AC_RUT_CARRERA_FALLBACK_KEY"] = composite.where(
        unmatched & composite.map(composite_counts).eq(1), ""
    )
    rut_table = tables.get("rut_carrera", pd.DataFrame()).copy()
    if not rut_table.empty:
        rut_table["_AC_RUT_CARRERA_FALLBACK_KEY"] = (
            rut_table["RUT_NORMALIZADO"].fillna("") + "|" + rut_table["CARRERA_NORMALIZADA"].fillna("")
        )
        rut_table = rut_table.drop(columns=["RUT_NORMALIZADO", "CARRERA_NORMALIZADA"])
        rename = {col: f"{col}__RUT" for col in EXPECTED_AC_COLS}
        rut_table.rename(columns=rename, inplace=True)
        rut_table["_AC_MATCH_RUT"] = True
    else:
        rut_table = pd.DataFrame(
            columns=["_AC_RUT_CARRERA_FALLBACK_KEY", "_AC_MATCH_RUT", *[f"{c}__RUT" for c in EXPECTED_AC_COLS]]
        )
    out = merge_preservando_filas(
        out,
        rut_table,
        on="_AC_RUT_CARRERA_FALLBACK_KEY",
        audit=audit,
        etapa="enriquecer_ac_rut_carrera_fallback",
        fuente="AC_REDUCIDA_RUT_CARRERA",
        validate="many_to_one",
    )
    apply_rut = out["FUENTE_MATCH_AC"].eq("") & out["_AC_MATCH_RUT"].astype("boolean").fillna(False)
    for col in EXPECTED_AC_COLS:
        candidate = f"{col}__RUT"
        out[col] = out[col].where(~apply_rut, out[candidate])
    out.loc[apply_rut, "FUENTE_MATCH_AC"] = "RUT_CARRERA_FALLBACK"
    fallback_count = int(out["FUENTE_MATCH_AC"].isin(["TOKEN_FALLBACK", "RUT_CARRERA_FALLBACK"]).sum())
    helper_cols = [
        "_AC_MATCH_CODCLI",
        "_AC_TOKEN_FALLBACK_KEY",
        "_AC_CARRERA_FUENTE",
        "_AC_MATCH_TOKEN",
        "_AC_RUT_CARRERA_FALLBACK_KEY",
        "_AC_MATCH_RUT",
        *[f"{col}__TOKEN" for col in EXPECTED_AC_COLS],
        *[f"{col}__RUT" for col in EXPECTED_AC_COLS],
    ]
    out.drop(columns=helper_cols, errors="ignore", inplace=True)
    return out, fallback_count


def clasificar_disponibilidad_ac(
    base: pd.DataFrame,
    *,
    columnas: Sequence[str] | None = None,
    minimo_variables_validas: int = 1,
) -> pd.DataFrame:
    columnas = list(columnas or AC_AVAILABILITY_COLS)
    if minimo_variables_validas < 1:
        raise ValueError("minimo_variables_validas debe ser al menos 1.")
    out = base.copy()
    available = [col for col in columnas if col in out.columns]
    valid_count = pd.Series(0, index=out.index, dtype="int64")
    for col in available:
        series = out[col]
        valid = series.notna()
        if not pd.api.types.is_numeric_dtype(series):
            valid &= series.map(_clean_identifier).ne("")
        valid_count = valid_count.add(valid.astype(int), fill_value=0).astype(int)
    out["CANTIDAD_VARIABLES_AC_VALIDAS"] = valid_count
    out["CON_DATOS_AC"] = valid_count.ge(int(minimo_variables_validas))
    out["ESTADO_DATOS_AC"] = np.where(out["CON_DATOS_AC"], "CON_DATOS_AC", "SIN_DATOS_AC")
    out["REGLA_DATOS_AC"] = (
        f"AL_MENOS_{int(minimo_variables_validas)}_DE_{len(columnas)}_VARIABLES_CONFIGURADAS"
    )
    return out


def _metricas_inscritos(
    inscritos: pd.DataFrame | None,
    *,
    token_col: str | None,
    key_kind: str,
) -> dict[str, int]:
    if inscritos is None or not isinstance(inscritos, pd.DataFrame) or inscritos.empty:
        return {"filas_inscritos": 0, "tokens_inscritos_unicos": 0, "tokens_inscritos_duplicados": 0}
    detected = _find_column(inscritos, token_col, ["token", "codigo de acceso", "access code"])
    if detected is None:
        raise SectorizadoValidationError("No se detectó la llave de inscritos LimeSurvey.")
    values = inscritos[detected].map(lambda value: normalize_cross_value(value, key_kind))
    valid = values[values.ne("")]
    return {
        "filas_inscritos": int(len(inscritos)),
        "tokens_inscritos_unicos": int(valid.nunique()),
        "tokens_inscritos_duplicados": int(valid.value_counts().gt(1).sum()),
    }


def construir_base_sectorizada_oficial(
    universo_vigente: pd.DataFrame,
    resultados: pd.DataFrame,
    *,
    anio: int,
    periodo: int,
    fecha_ejecucion: str,
    fuente_ac: pd.DataFrame | None = None,
    inscritos: pd.DataFrame | None = None,
    clave_resultados: str | None = None,
    clave_inscritos: str | None = None,
    tipo_llave_encuesta: str = "token",
    submitdate_col: str | None = None,
    datestamp_col: str | None = None,
    response_id_col: str | None = None,
    columnas_ac_requeridas: Sequence[str] | None = None,
    minimo_variables_ac: int = 1,
    execution_id: str | None = None,
    dedup_precomputado: DedupResult | None = None,
) -> SectorizadoResult:
    """Construye una fila por estudiante vigente y valida todos los invariantes."""
    execution_id = str(execution_id or fecha_ejecucion)
    audit = CardinalityAudit(execution_id, anio, periodo)
    base, metadata = construir_universo_maestro(
        universo_vigente,
        anio=anio,
        periodo=periodo,
        fecha_ejecucion=fecha_ejecucion,
        audit=audit,
        key_kind=tipo_llave_encuesta,
    )
    initial_rows = len(base)
    dedup = (
        dedup_precomputado
        if dedup_precomputado is not None
        else deduplicar_resultados_encuesta(
            resultados,
            token_col=clave_resultados,
            submitdate_col=submitdate_col,
            datestamp_col=datestamp_col,
            response_id_col=response_id_col,
            key_kind=tipo_llave_encuesta,
        )
    )
    if dedup_precomputado is not None and not dedup.resultados_unicos.equals(
        resultados
    ):
        raise SectorizadoValidationError(
            "El frame canónico recibido no coincide con la procedencia deduplicada."
        )
    audit.add(
        etapa="deduplicar_resultados_encuesta",
        fuente=(
            "RESULTADOS_LIMESURVEY_CANONICOS_CON_PROCEDENCIA_BRUTA"
            if dedup_precomputado is not None
            else "RESULTADOS_LIMESURVEY_BRUTOS"
        ),
        filas_antes=dedup.metricas["filas_resultados_brutos"],
        filas_despues=dedup.metricas["filas_resultados_unicos"],
        llave="_TOKEN_RESULTADO",
        key_df=dedup.resultados_unicos,
        cardinalidad_esperada="many_to_one_reduction",
        cardinalidad_observada="many_to_one_reduction",
        observacion=(
            "Se reutiliza la deduplicación institucional precomputada; métricas, "
            "selección y descartes conservan la fuente física original."
            if dedup_precomputado is not None
            else "Una fila seleccionada por token; descartes trazados en auditoría de duplicados."
        ),
    )
    survey_cols = [
        "_TOKEN_RESULTADO",
        "ESTADO_RESPUESTA_RESULTADO",
        "ID_RESPUESTA_SELECCIONADO",
        "FECHA_RESPUESTA_SELECCIONADA",
        "_SUBMITDATE_DT",
        "_DATESTAMP_DT",
    ]
    survey = dedup.resultados_unicos[[col for col in survey_cols if col in dedup.resultados_unicos.columns]].copy()
    token_counts_universe = base.loc[
        base["TOKEN_NORMALIZADO"].ne(""), "TOKEN_NORMALIZADO"
    ].value_counts()
    ambiguous_tokens = set(token_counts_universe[token_counts_universe.gt(1)].index)
    result_tokens = set(dedup.resultados_unicos["_TOKEN_RESULTADO"].dropna().astype(str))
    ambiguous_result_tokens = ambiguous_tokens & result_tokens
    base["TOKEN_AMBIGUO_UNIVERSO"] = base["TOKEN_NORMALIZADO"].isin(ambiguous_tokens)
    base["_TOKEN_CRUCE_RESULTADO"] = base["TOKEN_NORMALIZADO"].where(
        base["TOKEN_NORMALIZADO"].ne("") & ~base["TOKEN_AMBIGUO_UNIVERSO"],
        "",
    )
    audit.add(
        etapa="gobernar_llave_encuesta_univoca",
        fuente="DF_VIGENTES",
        filas_antes=initial_rows,
        filas_despues=len(base),
        llave="_TOKEN_CRUCE_RESULTADO",
        key_df=base,
        cardinalidad_esperada="one_to_one_en_llaves_no_nulas",
        cardinalidad_observada="one_to_one_en_llaves_no_nulas",
        observacion=(
            f"Tokens repetidos en universo={len(ambiguous_tokens)}; "
            f"tokens repetidos con resultado={len(ambiguous_result_tokens)}; "
            "no se atribuyen respuestas cuando el token no identifica un único estudiante vigente."
        ),
    )
    survey.rename(columns={"_TOKEN_RESULTADO": "_TOKEN_CRUCE_RESULTADO"}, inplace=True)
    base = merge_preservando_filas(
        base,
        survey,
        on="_TOKEN_CRUCE_RESULTADO",
        audit=audit,
        etapa="asociar_resultados_unicos",
        fuente="RESULTADOS_LIMESURVEY_DEDUPLICADOS",
        validate="many_to_one",
    )
    base["ESTADO_RESPUESTA"] = base["ESTADO_RESPUESTA_RESULTADO"].fillna("NO")
    invalid_states = ~base["ESTADO_RESPUESTA"].isin(RESPONSE_STATES)
    if bool(invalid_states.any()):
        raise SectorizadoValidationError("Se generaron estados de respuesta fuera del catálogo.", audit.frame())
    base["RESPONDE"] = base["ESTADO_RESPUESTA"]
    base["RESPONDE_BIN"] = base["ESTADO_RESPUESTA"].eq("SI").astype(int)
    base["RESPUESTA_PRESENTE_BIN"] = base["ESTADO_RESPUESTA"].isin(["SI", "INCOMPLETO"]).astype(int)
    base["MOTIVO_ESTADO_RESPUESTA"] = np.select(
        [
            base["ESTADO_RESPUESTA"].eq("SI"),
            base["ESTADO_RESPUESTA"].eq("INCOMPLETO"),
            base["TOKEN_AMBIGUO_UNIVERSO"]
            & base["TOKEN_NORMALIZADO"].isin(ambiguous_result_tokens),
            base["TOKEN_AMBIGUO_UNIVERSO"],
            base["TOKEN_NORMALIZADO"].eq(""),
        ],
        [
            "RESPUESTA_COMPLETA_ASOCIADA",
            "RESPUESTA_INCOMPLETA_ASOCIADA",
            "TOKEN_AMBIGUO_CON_RESULTADO_NO_ATRIBUIDO",
            "TOKEN_AMBIGUO_SIN_RESULTADO",
            "SIN_TOKEN_DE_ENCUESTA",
        ],
        default="SIN_RESULTADO_ASOCIADO",
    )

    ac_tables, conflicts_ac, ac_meta = reducir_fuente_ac(
        fuente_ac,
        universo=base,
        anio=anio,
        periodo=periodo,
        audit=audit,
    )
    base, ac_fallback_count = incorporar_ac(base, ac_tables, audit=audit)
    base = clasificar_disponibilidad_ac(
        base,
        columnas=columnas_ac_requeridas,
        minimo_variables_validas=minimo_variables_ac,
    )
    base["NUEVO_ANTIGUO"] = np.where(
        pd.to_numeric(base["ANOINGRESO"], errors="coerce").eq(int(anio)),
        "NUEVO",
        np.where(pd.to_numeric(base["ANOINGRESO"], errors="coerce").notna(), "ANTIGUO", NULL_LABEL),
    )
    for col in EXPECTED_AC_COLS[:-1]:
        base[col] = pd.to_numeric(base[col], errors="coerce")
    base["TASA_APROBACION_ANT"] = safe_div(base["_AC_ASI_APR_ANT"], base["_AC_ASI_INS_ANT"])
    base["TASA_APROBACION_HIS"] = safe_div(base["_AC_ASI_APR_HIS"], base["_AC_ASI_INS_HIS"])
    base["PROMEDIO_AC_REFERENCIA"] = base[["_AC_PROM_PRI_SEM", "_AC_PROM_SEG_SEM"]].mean(axis=1, skipna=True)
    base["TRAMO_PROMEDIO_AC"] = base["PROMEDIO_AC_REFERENCIA"].map(tramo_promedio)
    base["TRAMO_NIVEL_AC"] = base["_AC_NIV_ACA"].map(tramo_nivel)
    base["TRAMO_ANTIGUEDAD"] = base["ANOINGRESO"].map(lambda value: tramo_antiguedad(value, anio))

    inscritos_meta = _metricas_inscritos(
        inscritos,
        token_col=clave_inscritos,
        key_kind=tipo_llave_encuesta,
    )
    metadata.update(dedup.metricas)
    metadata.update(inscritos_meta)
    metadata.update(ac_meta)
    metadata["tokens_ambiguos_universo"] = int(len(ambiguous_tokens))
    metadata["filas_token_ambiguo_universo"] = int(base["TOKEN_AMBIGUO_UNIVERSO"].sum())
    metadata["tokens_ambiguos_con_resultado_no_atribuido"] = int(
        len(ambiguous_result_tokens)
    )
    metadata["filas_ambiguas_con_resultado_no_atribuido"] = int(
        (
            base["TOKEN_AMBIGUO_UNIVERSO"]
            & base["TOKEN_NORMALIZADO"].isin(ambiguous_result_tokens)
        ).sum()
    )
    metadata["casos_fallback_llave_ac"] = ac_fallback_count
    metadata["casos_con_fallback_llave"] = (
        metadata["casos_fallback_llave_universo"]
        + metadata["casos_token_desde_rut"]
        + ac_fallback_count
    )
    metadata["anio"] = int(anio)
    metadata["periodo"] = int(periodo)
    metadata["fecha_ejecucion"] = str(fecha_ejecucion)

    totals = base["ESTADO_RESPUESTA"].value_counts().reindex(RESPONSE_STATES, fill_value=0)
    si, incompleto, no = (int(totals[state]) for state in RESPONSE_STATES)
    con_ac = int(base["CON_DATOS_AC"].sum())
    sin_ac = int(len(base) - con_ac)
    present_mask = base["ESTADO_RESPUESTA"].isin(["SI", "INCOMPLETO"])
    associated_result_tokens = int(
        base.loc[present_mask, "_TOKEN_CRUCE_RESULTADO"].replace("", np.nan).nunique()
    )
    metadata["tokens_resultados_asociados_unicos"] = associated_result_tokens
    metadata["tokens_resultados_no_asociados"] = int(
        dedup.metricas["tokens_resultados_unicos"] - associated_result_tokens
    )
    errors: list[str] = []
    if len(base) != initial_rows:
        errors.append(f"base final {len(base)} != universo inicial {initial_rows}")
    if not base["LLAVE_ESTUDIANTE"].is_unique:
        errors.append("llave principal final no única")
    if si + incompleto + no != initial_rows:
        errors.append("estados no reconcilian")
    if con_ac + sin_ac != initial_rows:
        errors.append("disponibilidad AC no reconcilia")
    if si > dedup.metricas["tokens_resultados_unicos"]:
        errors.append("SI supera tokens únicos de resultados")
    if si + incompleto > dedup.metricas["tokens_resultados_unicos"]:
        errors.append("SI+INCOMPLETO supera tokens únicos de resultados")
    if si + incompleto != associated_result_tokens:
        errors.append("una respuesta única fue atribuida a más de un estudiante vigente")
    if errors:
        audit.add(
            etapa="validaciones_finales_bloqueantes",
            fuente="BASE_SECTORIZADA_FINAL",
            filas_antes=initial_rows,
            filas_despues=len(base),
            llave="LLAVE_ESTUDIANTE",
            key_df=base,
            cardinalidad_esperada="one_to_one",
            cardinalidad_observada="one_to_one" if base["LLAVE_ESTUDIANTE"].is_unique else "one_to_many",
            resultado="ERROR",
            observacion=" | ".join(errors),
        )
        raise SectorizadoValidationError("; ".join(errors), audit.frame())
    audit.add(
        etapa="validaciones_finales_bloqueantes",
        fuente="BASE_SECTORIZADA_FINAL",
        filas_antes=initial_rows,
        filas_despues=len(base),
        llave="LLAVE_ESTUDIANTE",
        key_df=base,
        cardinalidad_esperada="one_to_one",
        cardinalidad_observada="one_to_one",
        resultado="OK",
        observacion="Filas, llave, estados, AC y límites de resultados reconciliados.",
    )
    metadata.update(
        {
            "total_si": si,
            "total_incompleto": incompleto,
            "total_no": no,
            "total_con_datos_ac": con_ac,
            "total_sin_datos_ac": sin_ac,
        }
    )
    reconciliation = build_auditoria_reconciliacion_panel(base, metadata)
    return SectorizadoResult(
        base=base,
        auditoria_universo=audit.frame(),
        auditoria_duplicados=dedup.auditoria,
        auditoria_descartes_encuesta=dedup.descartes,
        auditoria_conflictos_ac=conflicts_ac,
        auditoria_reconciliacion_panel=reconciliation,
        metadata=metadata,
    )


def find_latest(pattern: str, base: Path) -> Path:
    files = sorted(base.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise SystemExit(f"{PY_PREFIX}[ERROR] No se encontraron archivos con patron {pattern} en {base}")
    return files[0]


def find_latest_encuesta_base(pattern: str, base: Path) -> Path:
    files = sorted(base.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    files = [p for p in files if "ANALISIS_SECTORIZADO" not in p.name and "AUDITORIA" not in p.name]
    if not files:
        raise SystemExit(f"{PY_PREFIX}[ERROR] No se encontro encuesta base con patron {pattern} en {base}")
    return files[0]


def find_latest_valid_trayectoria(pattern: str, base: Path) -> Path:
    files = sorted(base.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise SystemExit(f"{PY_PREFIX}[ERROR] No se encontraron archivos con patron {pattern} en {base}")

    for path in files:
        try:
            with pd.ExcelFile(path) as workbook:
                sheets = set(workbook.sheet_names)
            if {"TRAYECTORIA_POR_PERIODO", "AUDITORIA_AVANCE_CURRICULAR"} - sheets:
                continue
            cols = pd.read_excel(path, sheet_name="TRAYECTORIA_POR_PERIODO", nrows=0).columns
            missing_ac = [c for c in EXPECTED_AC_COLS if c not in cols]
            if missing_ac:
                continue
            return path
        except Exception as exc:
            log(f"[WARN] No se pudo validar trayectoria candidata {path}: {exc}")

    raise SystemExit(
        f"{PY_PREFIX}[ERROR] No se encontro trayectoria valida con TRAYECTORIA_POR_PERIODO, "
        "AUDITORIA_AVANCE_CURRICULAR y columnas _AC_."
    )


def parse_period_from_name(path: Path) -> tuple[int | None, int | None, str | None]:
    m = re.search(r"_(\d{4})_P(\d+)_(\d{8})", path.name)
    if not m:
        return None, None, None
    return int(m.group(1)), int(m.group(2)), m.group(3)


def default_output_path(encuesta: Path) -> Path:
    stem = encuesta.stem
    if stem.endswith("_ANALISIS_SECTORIZADO"):
        stem = stem[: -len("_ANALISIS_SECTORIZADO")]
    if stem.endswith("_ANALISIS_SECTORIZADO_OFICIAL"):
        stem = stem[: -len("_ANALISIS_SECTORIZADO_OFICIAL")]
    return encuesta.with_name(f"{stem}_ANALISIS_SECTORIZADO_OFICIAL.xlsx")


def classify_sector(carrera: Any) -> tuple[str, str]:
    txt = norm_txt(carrera)
    if not txt:
        return "Sin clasificar", "SIN_REGLA_CONSERVADORA"

    sector = _sector_oficial(txt)
    if sector:
        return sector, "MATCH_OFICIAL_EXACTO"

    if "CONTINUIDAD" in txt:
        for base in _bases_continuidad(txt):
            sector = _sector_oficial(base)
            if sector:
                return sector, "MATCH_OFICIAL_CONTINUIDAD"

    return "Sin clasificar", "SIN_REGLA_CONSERVADORA"


def _sector_oficial(carrera_norm: str) -> str | None:
    direct = CARRERA_A_SECTOR_OFICIAL.get(carrera_norm)
    if direct:
        return direct
    alias = ALIAS_CARRERA_OFICIAL.get(carrera_norm)
    if alias:
        return CARRERA_A_SECTOR_OFICIAL.get(alias)
    return None


def _bases_continuidad(carrera_norm: str) -> list[str]:
    base = re.sub(r"\bCONTINUIDAD\b", " ", carrera_norm)
    base = re.sub(r"\s+", " ", base).strip()
    if base.startswith("EN "):
        base = base[3:].strip()

    candidates = [base]
    if base.startswith("INGENIERIA ") and not base.startswith("INGENIERIA EN "):
        candidates.append(base.replace("INGENIERIA ", "INGENIERIA EN ", 1))
    if not base.startswith("TECNICO EN "):
        candidates.append(f"TECNICO EN {base}")
    if not base.startswith("INGENIERIA EN "):
        candidates.append(f"INGENIERIA EN {base}")

    out: list[str] = []
    for candidate in candidates:
        clean = norm_txt(candidate)
        if clean and clean not in out:
            out.append(clean)
    return out


def safe_div(num: pd.Series, den: pd.Series) -> pd.Series:
    num_n = pd.to_numeric(num, errors="coerce")
    den_n = pd.to_numeric(den, errors="coerce")
    values = np.where(den_n > 0, num_n / den_n, np.nan)
    return pd.Series(values, index=num.index, dtype="float64")


def tramo_promedio(value: Any) -> str:
    v = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(v):
        return "SIN_DATO"
    if v < 4.0:
        return "MENOR_4_0"
    if v < 5.0:
        return "4_0_A_4_9"
    if v < 6.0:
        return "5_0_A_5_9"
    return "6_0_O_MAS"


def tramo_nivel(value: Any) -> str:
    v = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(v):
        return "SIN_DATO"
    if v <= 2:
        return "NIVEL_1_2"
    if v <= 4:
        return "NIVEL_3_4"
    if v <= 6:
        return "NIVEL_5_6"
    return "NIVEL_7_O_MAS"


def tramo_antiguedad(anio_ingreso: Any, anio_ref: int | None) -> str:
    if anio_ref is None:
        return "SIN_DATO"
    anio = pd.to_numeric(pd.Series([anio_ingreso]), errors="coerce").iloc[0]
    if pd.isna(anio):
        return "SIN_DATO"
    delta = int(anio_ref) - int(anio)
    if delta <= 0:
        return "NUEVO"
    if delta == 1:
        return "1_ANIO"
    if delta == 2:
        return "2_ANIOS"
    return "3_O_MAS_ANIOS"


def read_structured_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=sheet_name, header=None, nrows=20)
    header_idx = None
    known = {
        "category_block",
        "category_name",
        "categoria_bloque",
        "categoria_nombre",
        "codigo_base",
        "codigo_columna",
        "question_text",
        "texto_pregunta",
        "answer",
        "pct_answered",
        "pct_si",
    }
    for idx, row in raw.iterrows():
        vals = {str(x).strip() for x in row.dropna().tolist()}
        if len(vals & known) >= 2:
            header_idx = int(str(idx))
            break
    if header_idx is None:
        return pd.DataFrame()

    df = pd.read_excel(path, sheet_name=sheet_name, header=header_idx)
    df = df.dropna(how="all").copy()
    df.columns = [str(c).strip() for c in df.columns]
    return df


def build_preguntas_clave(encuesta: Path) -> tuple[pd.DataFrame, list[str]]:
    warnings: list[str] = []
    frames: list[pd.DataFrame] = []
    keywords = [
        "ELECCION",
        "INSTITUCION",
        "CARRERA",
        "MEDIO",
        "INFORMACION",
        "MATRICULA",
        "BIENESTAR",
        "NEE",
        "DISCAPACIDAD",
        "SALUD",
        "APOYO",
        "TECNOLOG",
        "COMPUTADOR",
        "INTERNET",
        "EDUCACION SUPERIOR",
        "RED DE APOYO",
        "RECURSOS",
    ]
    for sheet in ["Analisis_respuestas", "Resumen_binario_ejecutivo"]:
        try:
            df = read_structured_sheet(encuesta, sheet)
        except Exception as exc:
            warnings.append(f"No fue posible leer {sheet}: {exc}")
            continue
        if df.empty:
            warnings.append(f"{sheet} no expone tabla estructurada legible.")
            continue

        text_cols = [
            c
            for c in df.columns
            if any(token in norm_txt(c) for token in ["PREGUNTA", "QUESTION", "CATEGORIA", "CATEGORY", "ANSWER", "RESPUESTA"])
        ]
        if not text_cols:
            text_cols = list(df.columns)
        haystack = df[text_cols].apply(
            lambda row: " ".join("" if pd.isna(x) else str(x) for x in row.tolist()),
            axis=1,
        ).map(norm_txt)
        mask = haystack.map(lambda x: any(k in str(x) for k in keywords))
        out = df.loc[mask].copy()
        if out.empty:
            warnings.append(f"{sheet} sin preguntas clave detectadas por reglas conservadoras.")
            continue
        out.insert(0, "FUENTE_HOJA", sheet)
        out.insert(
            1,
            "NOTA_METODOLOGICA",
            "Hoja agregada: no permite cruce individual por sector con el Excel de encuesta disponible.",
        )
        frames.append(out)

    if not frames:
        return (
            pd.DataFrame(
                [
                    {
                        "FUENTE_HOJA": "NO_DISPONIBLE",
                        "NOTA_METODOLOGICA": "No se detectaron preguntas clave en hojas agregadas.",
                    }
                ]
            ),
            warnings,
        )
    return pd.concat(frames, ignore_index=True).head(1000), warnings


def pct(series: pd.Series) -> float:
    return round(float(series.mean() * 100), 2) if len(series) else np.nan


def _fill_dimension(series: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(series):
        return series.astype("object").where(series.notna(), NULL_LABEL)
    cleaned = series.astype("object").where(series.notna(), NULL_LABEL)
    empty = cleaned.map(lambda value: _clean_identifier(value) == "")
    return cleaned.where(~empty, NULL_LABEL)


def agg_by(base: pd.DataFrame, dims: list[str]) -> pd.DataFrame:
    df = base.copy()
    for dim in dims:
        if dim not in df.columns:
            df[dim] = NULL_LABEL
        df[dim] = _fill_dimension(df[dim])
    if "ESTADO_RESPUESTA" not in df.columns:
        raise SectorizadoValidationError("La base final no contiene ESTADO_RESPUESTA.")
    df["_RESP_INCOMPLETA"] = df["ESTADO_RESPUESTA"].eq("INCOMPLETO").astype(int)
    df["_RESP_NO"] = df["ESTADO_RESPUESTA"].eq("NO").astype(int)
    grouped = (
        df.groupby(dims, dropna=False)
        .agg(
            ESTUDIANTES=("LLAVE_ESTUDIANTE", "size"),
            RESPONDIERON=("RESPONDE_BIN", "sum"),
            INCOMPLETOS=("_RESP_INCOMPLETA", "sum"),
            NO_RESPONDIERON=("_RESP_NO", "sum"),
            TASA_RESPUESTA_PCT=("RESPONDE_BIN", pct),
            TASA_INCLUYE_INCOMPLETOS_PCT=("RESPUESTA_PRESENTE_BIN", pct),
            PROMEDIO_AC=("PROMEDIO_AC_REFERENCIA", "mean"),
            NIVEL_AC_PROMEDIO=("_AC_NIV_ACA", "mean"),
            ASI_INS_HIS_PROMEDIO=("_AC_ASI_INS_HIS", "mean"),
            ASI_APR_HIS_PROMEDIO=("_AC_ASI_APR_HIS", "mean"),
            TASA_APROBACION_HIS=("TASA_APROBACION_HIS", "mean"),
        )
        .reset_index()
    )
    for col in ["PROMEDIO_AC", "NIVEL_AC_PROMEDIO", "ASI_INS_HIS_PROMEDIO", "ASI_APR_HIS_PROMEDIO", "TASA_APROBACION_HIS"]:
        if col in grouped.columns:
            grouped[col] = grouped[col].round(4)
    return grouped.sort_values(dims).reset_index(drop=True)


def construir_resumen_sectorizado(df: pd.DataFrame, group_cols: list[str], nombre_analisis: str) -> pd.DataFrame:
    base = df.copy()
    warnings: list[str] = []
    for col in group_cols:
        if col not in base.columns:
            base[col] = NULL_LABEL
            warnings.append(f"Columna no disponible para agrupacion: {col}")
        base[col] = _fill_dimension(base[col])

    for col in ["RESPONDE_BIN", "PROMEDIO_AC_REFERENCIA", "_AC_NIV_ACA", "TASA_APROBACION_ANT", "TASA_APROBACION_HIS"]:
        if col not in base.columns:
            base[col] = np.nan
    for col in ["_AC_ASI_INS_ANT", "_AC_ASI_APR_ANT", "_AC_ASI_INS_HIS", "_AC_ASI_APR_HIS"]:
        if col not in base.columns:
            base[col] = np.nan

    base["_CON_AC_RESUMEN"] = base.get("CON_DATOS_AC", pd.Series(False, index=base.index)).fillna(False).astype(bool)
    base["_RESP_INCOMPLETA"] = base.get("ESTADO_RESPUESTA", pd.Series("NO", index=base.index)).eq("INCOMPLETO").astype(int)
    base["_RESP_NO"] = base.get("ESTADO_RESPUESTA", pd.Series("NO", index=base.index)).eq("NO").astype(int)
    prom = pd.to_numeric(base["PROMEDIO_AC_REFERENCIA"], errors="coerce")
    base["_PROMEDIO_AC_VALIDO"] = prom.where(prom.between(1, 7))

    grouped = (
        base.groupby(group_cols, dropna=False)
        .agg(
            ESTUDIANTES=("LLAVE_ESTUDIANTE", "size"),
            RESPONDIERON=("RESPONDE_BIN", "sum"),
            INCOMPLETOS=("_RESP_INCOMPLETA", "sum"),
            NO_RESPONDIERON=("_RESP_NO", "sum"),
            CON_AC=("_CON_AC_RESUMEN", "sum"),
            PROMEDIO_AC_MEDIO=("_PROMEDIO_AC_VALIDO", "mean"),
            NIVEL_AC_PROMEDIO=("_AC_NIV_ACA", "mean"),
            ASI_INS_ANT_PROMEDIO=("_AC_ASI_INS_ANT", "mean"),
            ASI_APR_ANT_PROMEDIO=("_AC_ASI_APR_ANT", "mean"),
            ASI_INS_HIS_PROMEDIO=("_AC_ASI_INS_HIS", "mean"),
            ASI_APR_HIS_PROMEDIO=("_AC_ASI_APR_HIS", "mean"),
            TASA_APROBACION_ANT_MEDIA=("TASA_APROBACION_ANT", "mean"),
            TASA_APROBACION_HIS_MEDIA=("TASA_APROBACION_HIS", "mean"),
        )
        .reset_index()
    )
    grouped.insert(0, "NIVEL_ANALISIS", nombre_analisis)
    grouped["SIN_AC"] = grouped["ESTUDIANTES"] - grouped["CON_AC"]
    grouped["TASA_RESPUESTA_PCT"] = np.where(
        grouped["ESTUDIANTES"] > 0,
        (grouped["RESPONDIERON"] / grouped["ESTUDIANTES"] * 100).round(2),
        np.nan,
    )
    grouped["PORC_CON_AC"] = np.where(
        grouped["ESTUDIANTES"] > 0,
        (grouped["CON_AC"] / grouped["ESTUDIANTES"] * 100).round(2),
        np.nan,
    )
    grouped["TASA_INCLUYE_INCOMPLETOS_PCT"] = np.where(
        grouped["ESTUDIANTES"] > 0,
        ((grouped["RESPONDIERON"] + grouped["INCOMPLETOS"]) / grouped["ESTUDIANTES"] * 100).round(2),
        np.nan,
    )

    numeric_cols = [
        "PROMEDIO_AC_MEDIO",
        "NIVEL_AC_PROMEDIO",
        "ASI_INS_ANT_PROMEDIO",
        "ASI_APR_ANT_PROMEDIO",
        "ASI_INS_HIS_PROMEDIO",
        "ASI_APR_HIS_PROMEDIO",
        "TASA_APROBACION_ANT_MEDIA",
        "TASA_APROBACION_HIS_MEDIA",
    ]
    for col in numeric_cols:
        grouped[col] = pd.to_numeric(grouped[col], errors="coerce").round(4)

    preferred = [
        "NIVEL_ANALISIS",
        "SECTOR_ACADEMICO",
        "JORNADA",
        "NUEVO_ANTIGUO",
        "CARRERA",
        "CODCARPR",
        "ANOINGRESO",
        "TRAMO_NIVEL_AC",
        "TRAMO_PROMEDIO_AC",
        "ESTUDIANTES",
        "RESPONDIERON",
        "INCOMPLETOS",
        "NO_RESPONDIERON",
        "TASA_RESPUESTA_PCT",
        "TASA_INCLUYE_INCOMPLETOS_PCT",
        "CON_AC",
        "SIN_AC",
        "PORC_CON_AC",
        "PROMEDIO_AC_MEDIO",
        "NIVEL_AC_PROMEDIO",
        "ASI_INS_ANT_PROMEDIO",
        "ASI_APR_ANT_PROMEDIO",
        "ASI_INS_HIS_PROMEDIO",
        "ASI_APR_HIS_PROMEDIO",
        "TASA_APROBACION_ANT_MEDIA",
        "TASA_APROBACION_HIS_MEDIA",
    ]
    ordered = [c for c in preferred if c in grouped.columns] + [c for c in grouped.columns if c not in preferred]
    grouped = grouped[ordered].sort_values([c for c in group_cols if c in grouped.columns]).reset_index(drop=True)
    if warnings:
        grouped["WARN_AGRUPACION"] = " | ".join(warnings)
    return grouped


def build_matriz_sector_jornada_na(base: pd.DataFrame) -> pd.DataFrame:
    resumen = construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO", "JORNADA", "NUEVO_ANTIGUO"], "MATRIZ_SECTOR_JORNADA_NA")
    rows: list[dict[str, Any]] = []
    for _, row in resumen.iterrows():
        segmento = f"{row.get('JORNADA', 'SIN_DATO')} | {row.get('NUEVO_ANTIGUO', 'SIN_DATO')}"
        rows.extend(
            [
                {
                    "SECTOR_ACADEMICO": row["SECTOR_ACADEMICO"],
                    "SEGMENTO_JORNADA_NA": segmento,
                    "INDICADOR": "ESTUDIANTES",
                    "VALOR": row["ESTUDIANTES"],
                },
                {
                    "SECTOR_ACADEMICO": row["SECTOR_ACADEMICO"],
                    "SEGMENTO_JORNADA_NA": segmento,
                    "INDICADOR": "RESPONDIERON",
                    "VALOR": row["RESPONDIERON"],
                },
                {
                    "SECTOR_ACADEMICO": row["SECTOR_ACADEMICO"],
                    "SEGMENTO_JORNADA_NA": segmento,
                    "INDICADOR": "TASA_RESPUESTA_PCT",
                    "VALOR": row["TASA_RESPUESTA_PCT"],
                },
            ]
        )
    return pd.DataFrame(rows)


def build_lectura_ejecutiva_sector(base: pd.DataFrame) -> pd.DataFrame:
    sector = construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO"], "SECTOR")
    jornada = construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO", "JORNADA"], "SECTOR_JORNADA")
    nuevo = construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO", "NUEVO_ANTIGUO"], "SECTOR_NUEVO_ANTIGUO")
    carrera_seg = construir_resumen_sectorizado(
        base, ["SECTOR_ACADEMICO", "CARRERA", "JORNADA", "NUEVO_ANTIGUO"], "CARRERA_JORNADA_NA"
    )
    rows: list[dict[str, Any]] = []

    if not sector.empty:
        max_sector = sector.sort_values("TASA_RESPUESTA_PCT", ascending=False).iloc[0]
        min_sector = sector.sort_values("TASA_RESPUESTA_PCT", ascending=True).iloc[0]
        rows.append(
            {
                "LECTURA": "SECTOR_MAYOR_RESPUESTA",
                "SECTOR_ACADEMICO": max_sector["SECTOR_ACADEMICO"],
                "SEGMENTO": "",
                "VALOR": max_sector["TASA_RESPUESTA_PCT"],
                "DETALLE": f"Estudiantes={int(max_sector['ESTUDIANTES'])}; respondieron={int(max_sector['RESPONDIERON'])}",
            }
        )
        rows.append(
            {
                "LECTURA": "SECTOR_MENOR_RESPUESTA",
                "SECTOR_ACADEMICO": min_sector["SECTOR_ACADEMICO"],
                "SEGMENTO": "",
                "VALOR": min_sector["TASA_RESPUESTA_PCT"],
                "DETALLE": f"Estudiantes={int(min_sector['ESTUDIANTES'])}; respondieron={int(min_sector['RESPONDIERON'])}",
            }
        )

    for sector_name, group in jornada.groupby("SECTOR_ACADEMICO", dropna=False):
        if group.empty:
            continue
        row = group.sort_values("TASA_RESPUESTA_PCT", ascending=True).iloc[0]
        rows.append(
            {
                "LECTURA": "JORNADA_MENOR_TASA_EN_SECTOR",
                "SECTOR_ACADEMICO": sector_name,
                "SEGMENTO": row["JORNADA"],
                "VALOR": row["TASA_RESPUESTA_PCT"],
                "DETALLE": f"Estudiantes={int(row['ESTUDIANTES'])}; respondieron={int(row['RESPONDIERON'])}",
            }
        )

    for sector_name, group in nuevo.groupby("SECTOR_ACADEMICO", dropna=False):
        pivot = group.set_index("NUEVO_ANTIGUO")["TASA_RESPUESTA_PCT"].to_dict()
        if "NUEVO" in pivot and "ANTIGUO" in pivot:
            rows.append(
                {
                    "LECTURA": "BRECHA_NUEVO_ANTIGUO",
                    "SECTOR_ACADEMICO": sector_name,
                    "SEGMENTO": "NUEVO - ANTIGUO",
                    "VALOR": round(float(pivot["NUEVO"] - pivot["ANTIGUO"]), 2),
                    "DETALLE": f"NUEVO={pivot['NUEVO']}; ANTIGUO={pivot['ANTIGUO']}",
                }
            )

    criticos = carrera_seg[(carrera_seg["ESTUDIANTES"] >= 10) & (carrera_seg["TASA_RESPUESTA_PCT"] < 40)].copy()
    for _, row in criticos.sort_values(["TASA_RESPUESTA_PCT", "ESTUDIANTES"], ascending=[True, False]).head(25).iterrows():
        rows.append(
            {
                "LECTURA": "SEGMENTO_CRITICO_RESPUESTA",
                "SECTOR_ACADEMICO": row["SECTOR_ACADEMICO"],
                "SEGMENTO": f"{row['CARRERA']} | {row['JORNADA']} | {row['NUEVO_ANTIGUO']}",
                "VALOR": row["TASA_RESPUESTA_PCT"],
                "DETALLE": f"Estudiantes={int(row['ESTUDIANTES'])}; respondieron={int(row['RESPONDIERON'])}. Interpretar con cautela segmentos pequenos.",
            }
        )

    rows.append(
        {
            "LECTURA": "ADVERTENCIA_METODOLOGICA",
            "SECTOR_ACADEMICO": "",
            "SEGMENTO": "",
            "VALOR": "",
            "DETALLE": "Segmentos con pocos casos pueden cambiar mucho ante variaciones pequenas; priorizar segmentos con ESTUDIANTES >= 10.",
        }
    )
    return pd.DataFrame(rows)


def build_auditoria_granularidad(sheets: dict[str, pd.DataFrame], base: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    base_rows = len(base)
    for sheet in GRANULARITY_SHEETS:
        if sheet == "AUDITORIA_GRANULARIDAD":
            continue
        df = sheets.get(sheet, pd.DataFrame())
        if "ESTUDIANTES" in df.columns:
            reconciled: int | str = int(pd.to_numeric(df["ESTUDIANTES"], errors="coerce").fillna(0).sum())
            result = "OK" if reconciled == base_rows else "ERROR"
        else:
            reconciled = "NO_APLICA_HOJA_NO_PARTICIONAL"
            result = "OK"
        rows.append(
            {
                "HOJA": sheet,
                "FILAS": len(df),
                "CONTIENE_SECTOR_ACADEMICO": "SECTOR_ACADEMICO" in df.columns,
                "CONTIENE_JORNADA": "JORNADA" in df.columns,
                "CONTIENE_NUEVO_ANTIGUO": "NUEVO_ANTIGUO" in df.columns,
                "ESTUDIANTES_RECONCILIADOS": reconciled,
                "BASE_SECTORIZADA_FILAS": base_rows,
                "RESULTADO": result,
            }
        )
        if result == "ERROR":
            raise SectorizadoValidationError(
                f"La hoja {sheet} no reconcilia con BASE_SECTORIZADA: {reconciled} != {base_rows}."
            )
    rows.append(
        {
            "HOJA": "NOTA",
            "FILAS": "",
            "CONTIENE_SECTOR_ACADEMICO": "",
            "CONTIENE_JORNADA": "",
            "CONTIENE_NUEVO_ANTIGUO": "",
            "ESTUDIANTES_RECONCILIADOS": "",
            "BASE_SECTORIZADA_FILAS": base_rows,
            "RESULTADO": "El detalle carrera-sector esta en BASE_SECTORIZADA, SECTOR_CARRERA y AUDITORIA_SECTORES.",
        }
    )
    return pd.DataFrame(rows)


def build_auditoria_sectores(base: pd.DataFrame) -> pd.DataFrame:
    tmp = base.copy()
    tmp["CARRERA_NORMALIZADA"] = tmp["CARRERA"].map(norm_txt)
    for col in ["CARRERA", "CARRERA_NORMALIZADA", "SECTOR_ACADEMICO", "REGLA_SECTOR_ACADEMICO"]:
        tmp[col] = _fill_dimension(tmp[col])
    out = (
        tmp.groupby(["CARRERA", "CARRERA_NORMALIZADA", "SECTOR_ACADEMICO", "REGLA_SECTOR_ACADEMICO"], dropna=False)
        .agg(
            ESTUDIANTES=("LLAVE_ESTUDIANTE", "size"),
            RESPONDIERON=("RESPONDE_BIN", "sum"),
            INCOMPLETOS=("ESTADO_RESPUESTA", lambda s: int(s.eq("INCOMPLETO").sum())),
        )
        .reset_index()
    )
    out["TASA_RESPUESTA_PCT"] = np.where(out["ESTUDIANTES"] > 0, (out["RESPONDIERON"] / out["ESTUDIANTES"] * 100).round(2), np.nan)
    return out.sort_values(["SECTOR_ACADEMICO", "ESTUDIANTES"], ascending=[True, False]).reset_index(drop=True)


def build_panel_kpis(base: pd.DataFrame, metadata: dict[str, Any]) -> pd.DataFrame:
    total = int(len(base))
    si = int(base["ESTADO_RESPUESTA"].eq("SI").sum())
    incompleto = int(base["ESTADO_RESPUESTA"].eq("INCOMPLETO").sum())
    no = int(base["ESTADO_RESPUESTA"].eq("NO").sum())
    con_ac = int(base["CON_DATOS_AC"].sum())

    def distinct_valid(col: str) -> int:
        if col not in base.columns:
            return 0
        values = base[col].map(_clean_identifier)
        return int(values[values.ne("") & values.ne(NULL_LABEL)].nunique())

    rows = [
        ("Año de análisis", int(metadata["anio"]), "ENTERO"),
        ("Período de análisis", int(metadata["periodo"]), "ENTERO"),
        ("Fecha de ejecución", metadata["fecha_ejecucion"], "FECHA"),
        ("Total estudiantes vigentes", total, "ENTERO"),
        ("Respondieron completamente", si, "ENTERO"),
        ("Respuesta incompleta", incompleto, "ENTERO"),
        ("No respondieron", no, "ENTERO"),
        ("Tasa de respuesta completa", si / total if total else np.nan, "PORCENTAJE"),
        ("Tasa incluyendo incompletos", (si + incompleto) / total if total else np.nan, "PORCENTAJE"),
        ("Total sectores", distinct_valid("SECTOR_ACADEMICO"), "ENTERO"),
        ("Total carreras", distinct_valid("CARRERA"), "ENTERO"),
        ("Total jornadas", distinct_valid("JORNADA"), "ENTERO"),
        ("Total sedes", distinct_valid("SEDE"), "ENTERO"),
        ("Total con datos AC", con_ac, "ENTERO"),
        ("Total sin datos AC", total - con_ac, "ENTERO"),
        ("Tokens únicos en resultados", int(metadata.get("tokens_resultados_unicos", 0)), "ENTERO"),
        (
            "Tokens duplicados en resultados brutos",
            int(metadata.get("tokens_duplicados_resultados_brutos", 0)),
            "ENTERO",
        ),
        ("Filas asociadas a duplicados", int(metadata.get("filas_asociadas_duplicados", 0)), "ENTERO"),
        (
            "Tokens ambiguos con resultado no atribuido",
            int(metadata.get("tokens_ambiguos_con_resultado_no_atribuido", 0)),
            "ENTERO",
        ),
        (
            "Tokens de resultados no asociados al universo",
            int(metadata.get("tokens_resultados_no_asociados", 0)),
            "ENTERO",
        ),
        ("Conflictos académicos pendientes", int(metadata.get("conflictos_ac_pendientes", 0)), "ENTERO"),
        ("Casos con fallback de llave", int(metadata.get("casos_con_fallback_llave", 0)), "ENTERO"),
    ]
    return pd.DataFrame(rows, columns=["KPI", "VALOR", "FORMATO"])


def build_auditoria_reconciliacion_panel(base: pd.DataFrame, metadata: dict[str, Any]) -> pd.DataFrame:
    total = int(len(base))
    si = int(base.get("ESTADO_RESPUESTA", pd.Series(dtype="object")).eq("SI").sum())
    incompleto = int(base.get("ESTADO_RESPUESTA", pd.Series(dtype="object")).eq("INCOMPLETO").sum())
    no = int(base.get("ESTADO_RESPUESTA", pd.Series(dtype="object")).eq("NO").sum())
    con_ac = int(base.get("CON_DATOS_AC", pd.Series(False, index=base.index)).fillna(False).sum())
    comparisons = [
        ("Total vigentes", int(metadata.get("total_vigentes", total)), total),
        ("SI", si, int(metadata.get("total_si", si))),
        ("INCOMPLETO", incompleto, int(metadata.get("total_incompleto", incompleto))),
        ("NO", no, int(metadata.get("total_no", no))),
        ("Total estados", total, si + incompleto + no),
        ("Con datos AC", con_ac, int(metadata.get("total_con_datos_ac", con_ac))),
        ("Sin datos AC", total - con_ac, int(metadata.get("total_sin_datos_ac", total - con_ac))),
    ]
    rows = []
    for metric, official, panel in comparisons:
        difference = int(panel - official)
        rows.append(
            {
                "Métrica": metric,
                "Fuente oficial": int(official),
                "Panel sectorizado": int(panel),
                "Diferencia": difference,
                "Resultado": "OK" if difference == 0 else "ERROR",
            }
        )
    return pd.DataFrame(rows)


def _append_group_audit(
    result: SectorizadoResult,
    *,
    sheet_name: str,
    grouped: pd.DataFrame,
    total_column: str,
) -> None:
    expected = int(len(result.base))
    observed = int(pd.to_numeric(grouped.get(total_column, pd.Series(dtype="float64")), errors="coerce").fillna(0).sum())
    ok = expected == observed
    row = {
        "ejecución": str(result.metadata.get("fecha_ejecucion", "")),
        "año": int(result.metadata["anio"]),
        "período": int(result.metadata["periodo"]),
        "etapa": f"reconciliar_agrupacion_{sheet_name}",
        "fuente": "BASE_SECTORIZADA_FINAL",
        "filas_antes": expected,
        "filas_despues": expected,
        "diferencia": 0,
        "llave": "LLAVE_ESTUDIANTE",
        "llaves_unicas": expected,
        "llaves_nulas": 0,
        "llaves_duplicadas": 0,
        "cardinalidad_esperada": f"suma_{total_column}={expected}",
        "cardinalidad_observada": f"suma_{total_column}={observed}",
        "resultado": "OK" if ok else "ERROR",
        "observación": "Agrupación derivada exclusivamente de la base final con nulos explícitos.",
    }
    result.auditoria_universo = pd.concat(
        [result.auditoria_universo, pd.DataFrame([row], columns=AUDIT_UNIVERSE_COLUMNS)],
        ignore_index=True,
    )
    if not ok:
        raise SectorizadoValidationError(
            f"La agrupación {sheet_name} no reconcilia: {observed} != {expected}.",
            result.auditoria_universo,
        )


def build_outputs(encuesta: Path, result: SectorizadoResult) -> dict[str, pd.DataFrame]:
    base = result.base.copy()
    base["FUENTE_CRUCE"] = "DF_VIGENTES"
    sectores = base["CARRERA"].map(classify_sector)
    base["SECTOR_ACADEMICO"] = sectores.map(
        lambda item: item[0] if isinstance(item, tuple) else "Sin clasificar"
    )
    base["REGLA_SECTOR_ACADEMICO"] = sectores.map(
        lambda item: item[1] if isinstance(item, tuple) else "SIN_REGLA_CONSERVADORA"
    )
    result.base = base
    result.auditoria_reconciliacion_panel = build_auditoria_reconciliacion_panel(base, result.metadata)
    preguntas, preguntas_warn = build_preguntas_clave(encuesta)
    preguntas["CLASIFICACION_GRANULARIDAD"] = "AGREGADO_GLOBAL"
    preguntas["ESTADO_SECTORIZACION"] = "NO_SECTORIZABLE_CON_FUENTE_ACTUAL"
    auditoria_sectores = build_auditoria_sectores(base)

    total = int(len(base))
    respondieron = int(base["ESTADO_RESPUESTA"].eq("SI").sum())
    incompletos = int(base["ESTADO_RESPUESTA"].eq("INCOMPLETO").sum())
    no_resp = int(base["ESTADO_RESPUESTA"].eq("NO").sum())
    sector_resp = agg_by(base, ["SECTOR_ACADEMICO"])
    sector_valid = sector_resp[sector_resp["ESTUDIANTES"] > 0].copy()
    max_sector = sector_valid.sort_values("TASA_RESPUESTA_PCT", ascending=False).head(1)
    min_sector = sector_valid.sort_values("TASA_RESPUESTA_PCT", ascending=True).head(1)

    resumen_rows = [
        {"SECCION": "GENERAL", "INDICADOR": "TOTAL_ESTUDIANTES", "SECTOR": "", "VALOR": total, "DETALLE": ""},
        {"SECCION": "GENERAL", "INDICADOR": "TOTAL_RESPONDIERON", "SECTOR": "", "VALOR": respondieron, "DETALLE": ""},
        {"SECCION": "GENERAL", "INDICADOR": "TOTAL_INCOMPLETOS", "SECTOR": "", "VALOR": incompletos, "DETALLE": ""},
        {"SECCION": "GENERAL", "INDICADOR": "TOTAL_NO_RESPONDIERON", "SECTOR": "", "VALOR": no_resp, "DETALLE": ""},
        {"SECCION": "GENERAL", "INDICADOR": "TASA_RESPUESTA", "SECTOR": "", "VALOR": respondieron / total if total else np.nan, "DETALLE": "Valor numérico; formato porcentaje."},
        {
            "SECCION": "OBSERVACION",
            "INDICADOR": "DECISION_METODOLOGICA",
            "SECTOR": "",
            "VALOR": "",
            "DETALLE": "Base individual tomada exclusivamente del universo oficial vigente; la encuesta agregada no se atribuye a estudiantes.",
        },
    ]
    if not max_sector.empty:
        resumen_rows.append(
            {
                "SECCION": "SECTOR",
                "INDICADOR": "SECTOR_MAYOR_RESPUESTA",
                "SECTOR": max_sector.iloc[0]["SECTOR_ACADEMICO"],
                "VALOR": max_sector.iloc[0]["TASA_RESPUESTA_PCT"],
                "DETALLE": "",
            }
        )
    if not min_sector.empty:
        resumen_rows.append(
            {
                "SECCION": "SECTOR",
                "INDICADOR": "SECTOR_MENOR_RESPUESTA",
                "SECTOR": min_sector.iloc[0]["SECTOR_ACADEMICO"],
                "VALOR": min_sector.iloc[0]["TASA_RESPUESTA_PCT"],
                "DETALLE": "",
            }
        )
    for _, row in sector_resp.iterrows():
        resumen_rows.append(
            {
                "SECCION": "DISTRIBUCION_SECTOR",
                "INDICADOR": "ESTUDIANTES",
                "SECTOR": row["SECTOR_ACADEMICO"],
                "VALOR": int(row["ESTUDIANTES"]),
                "DETALLE": f"Tasa respuesta={row['TASA_RESPUESTA_PCT']}",
            }
        )
        resumen_rows.append(
            {
                "SECCION": "DISTRIBUCION_SECTOR",
                "INDICADOR": "PROMEDIO_AC",
                "SECTOR": row["SECTOR_ACADEMICO"],
                "VALOR": row["PROMEDIO_AC"],
                "DETALLE": f"Nivel AC promedio={row['NIVEL_AC_PROMEDIO']}",
            }
        )

    response_base = base.copy()
    response_base["SECTOR_ACADEMICO"] = _fill_dimension(response_base["SECTOR_ACADEMICO"])
    sector_respuesta = (
        response_base.groupby(["SECTOR_ACADEMICO", "ESTADO_RESPUESTA"], dropna=False)
        .agg(TOTAL=("LLAVE_ESTUDIANTE", "size"))
        .reset_index()
    )
    sector_respuesta.rename(columns={"ESTADO_RESPUESTA": "RESPONDE"}, inplace=True)
    sector_totals = sector_respuesta.groupby("SECTOR_ACADEMICO")["TOTAL"].transform("sum")
    sector_respuesta["PORC_DENTRO_SECTOR"] = np.where(sector_totals > 0, (sector_respuesta["TOTAL"] / sector_totals * 100).round(2), np.nan)
    sector_respuesta["PORC_GLOBAL"] = np.where(total > 0, (sector_respuesta["TOTAL"] / total * 100).round(2), np.nan)
    sector_respuesta = sector_respuesta.merge(
        sector_resp[["SECTOR_ACADEMICO", "ESTUDIANTES", "RESPONDIERON", "TASA_RESPUESTA_PCT"]],
        on="SECTOR_ACADEMICO",
        how="left",
        validate="many_to_one",
    )
    if int(sector_respuesta["TOTAL"].sum()) != total:
        raise SectorizadoValidationError("SECTOR_RESPUESTA no reconcilia con el universo.", result.auditoria_universo)

    auditoria_rows = [
        {"CAMPO": "archivo_encuesta_usado", "VALOR": str(encuesta)},
        {"CAMPO": "fuente_universo", "VALOR": "DF_VIGENTES"},
        {"CAMPO": "fecha_hora_generacion", "VALOR": result.metadata["fecha_ejecucion"]},
        {"CAMPO": "anio_analisis", "VALOR": result.metadata["anio"]},
        {"CAMPO": "periodo_analisis", "VALOR": result.metadata["periodo"]},
        {"CAMPO": "diccionario_sectorial", "VALOR": "OFICIAL_IMAGEN_OFERTA_ACADEMICA_FASE_15B"},
        {"CAMPO": "diccionario_sectorial_actualizacion", "VALOR": "FASE_16B_TECNICO_EN_PROGRAMACION_COMPUTACIONAL_TECNOLOGIA"},
        {"CAMPO": "regla_continuidad", "VALOR": REGLA_CONTINUIDAD_OFICIAL["descripcion"]},
        {"CAMPO": "llave_usada", "VALOR": "CODCLI > TOKEN > RUT con DV > llave técnica compuesta; fallbacks auditados."},
        {"CAMPO": "filas_encuesta_individuales", "VALOR": result.metadata["filas_resultados_unicos"]},
        {"CAMPO": "filas_base_sectorizada", "VALOR": total},
        {"CAMPO": "filas_cruzadas", "VALOR": total},
        {"CAMPO": "filas_sin_match", "VALOR": no_resp},
        {"CAMPO": "columnas_ac_usadas", "VALOR": ", ".join(result.metadata.get("columnas_ac_disponibles", []))},
        {"CAMPO": "sectores_definidos", "VALOR": ", ".join(sorted(base["SECTOR_ACADEMICO"].dropna().unique()))},
        {"CAMPO": "decision_metodologica", "VALOR": "El sectorizado desagrega el snapshot oficial vigente y no reconstruye su población."},
        {"CAMPO": "limitacion_preguntas_clave", "VALOR": "Analisis de preguntas clave queda agregado; no se sectoriza por pregunta al no existir base individual en el Excel encuesta."},
        {"CAMPO": "trazabilidad_carrera_sector", "VALOR": "La trazabilidad por carrera esta en BASE_SECTORIZADA, SECTOR_CARRERA y AUDITORIA_SECTORES."},
        {"CAMPO": "trazabilidad_hojas_agregadas", "VALOR": "Las hojas agregadas por sector/anio/nivel/promedio conservan SECTOR_ACADEMICO, ESTUDIANTES, RESPONDIERON y TASA_RESPUESTA_PCT cuando aplica."},
        {"CAMPO": "granularidad_jornada", "VALOR": "SI"},
        {"CAMPO": "granularidad_nuevo_antiguo", "VALOR": "SI"},
        {"CAMPO": "hojas_granularidad_agregadas", "VALOR": ", ".join(GRANULARITY_SHEETS)},
        {"CAMPO": "diccionario_sectorial_oficial", "VALOR": "SI"},
        {"CAMPO": "sin_clasificar", "VALOR": int(base["SECTOR_ACADEMICO"].eq("Sin clasificar").sum())},
    ]
    for key, value in sorted(result.metadata.items()):
        if key not in {"columnas_ac_disponibles"}:
            auditoria_rows.append({"CAMPO": key, "VALOR": value})
    for warn in preguntas_warn:
        auditoria_rows.append({"CAMPO": "advertencia", "VALOR": warn})

    outputs = {
        "SECTOR_RESUMEN_EJECUTIVO": pd.DataFrame(resumen_rows),
        "SECTOR_RESPUESTA": sector_respuesta.sort_values(["SECTOR_ACADEMICO", "RESPONDE"]).reset_index(drop=True),
        "SECTOR_CARRERA": agg_by(base, ["SECTOR_ACADEMICO", "CARRERA", "CARRERA_NORMALIZADA", "CODCARPR", "REGLA_SECTOR_ACADEMICO"]),
        "SECTOR_ANO_INGRESO": agg_by(base, ["SECTOR_ACADEMICO", "ANOINGRESO"]),
        "SECTOR_NUEVO_ANTIGUO": agg_by(base, ["SECTOR_ACADEMICO", "NUEVO_ANTIGUO"]),
        "SECTOR_NIVEL_AC": agg_by(base, ["SECTOR_ACADEMICO", "TRAMO_NIVEL_AC"]),
        "SECTOR_PROMEDIOS_AC": agg_by(base, ["SECTOR_ACADEMICO", "TRAMO_PROMEDIO_AC"]),
        "SECTOR_PREGUNTAS_CLAVE": preguntas,
        "BASE_SECTORIZADA": base,
        "AUDITORIA_SECTORES": auditoria_sectores,
        "AUDITORIA_CRUCE_AC": pd.DataFrame(auditoria_rows),
        "AUDITORIA_UNIVERSO_SECTORIZADO": result.auditoria_universo,
        "AUDITORIA_DUPLICADOS_ENCUESTA": result.auditoria_duplicados,
        "AUDITORIA_DESCARTES_ENCUESTA": result.auditoria_descartes_encuesta,
        "AUDITORIA_CONFLICTOS_AC": result.auditoria_conflictos_ac,
        "AUDITORIA_RECONCILIACION_PANEL": result.auditoria_reconciliacion_panel,
    }
    outputs.update(
        {
            "SECTOR_JORNADA": construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO", "JORNADA"], "SECTOR_JORNADA"),
            "SECTOR_JORNADA_NA": construir_resumen_sectorizado(
                base, ["SECTOR_ACADEMICO", "JORNADA", "NUEVO_ANTIGUO"], "SECTOR_JORNADA_NA"
            ),
            "CARRERA_JORNADA_NA": construir_resumen_sectorizado(
                base, ["SECTOR_ACADEMICO", "CARRERA", "CODCARPR", "JORNADA", "NUEVO_ANTIGUO"], "CARRERA_JORNADA_NA"
            ),
            "ANO_INGRESO_JORNADA_NA": construir_resumen_sectorizado(
                base, ["SECTOR_ACADEMICO", "ANOINGRESO", "JORNADA", "NUEVO_ANTIGUO"], "ANO_INGRESO_JORNADA_NA"
            ),
            "NIVEL_AC_JORNADA_NA": construir_resumen_sectorizado(
                base, ["SECTOR_ACADEMICO", "TRAMO_NIVEL_AC", "JORNADA", "NUEVO_ANTIGUO"], "NIVEL_AC_JORNADA_NA"
            ),
            "PROMEDIO_AC_JORNADA_NA": construir_resumen_sectorizado(
                base, ["SECTOR_ACADEMICO", "TRAMO_PROMEDIO_AC", "JORNADA", "NUEVO_ANTIGUO"], "PROMEDIO_AC_JORNADA_NA"
            ),
            "MATRIZ_SECTOR_JORNADA_NA": build_matriz_sector_jornada_na(base),
            "LECTURA_EJECUTIVA_SECTOR": build_lectura_ejecutiva_sector(base),
        }
    )
    partition_sheets = {
        "SECTOR_RESPUESTA": "TOTAL",
        "SECTOR_CARRERA": "ESTUDIANTES",
        "SECTOR_ANO_INGRESO": "ESTUDIANTES",
        "SECTOR_NUEVO_ANTIGUO": "ESTUDIANTES",
        "SECTOR_NIVEL_AC": "ESTUDIANTES",
        "SECTOR_PROMEDIOS_AC": "ESTUDIANTES",
        "SECTOR_JORNADA": "ESTUDIANTES",
        "SECTOR_JORNADA_NA": "ESTUDIANTES",
        "CARRERA_JORNADA_NA": "ESTUDIANTES",
        "ANO_INGRESO_JORNADA_NA": "ESTUDIANTES",
        "NIVEL_AC_JORNADA_NA": "ESTUDIANTES",
        "PROMEDIO_AC_JORNADA_NA": "ESTUDIANTES",
    }
    for sheet_name, total_col in partition_sheets.items():
        _append_group_audit(result, sheet_name=sheet_name, grouped=outputs[sheet_name], total_column=total_col)
    outputs["AUDITORIA_UNIVERSO_SECTORIZADO"] = result.auditoria_universo
    outputs["AUDITORIA_GRANULARIDAD"] = build_auditoria_granularidad(outputs, base)
    return outputs


def _format_sheet(ws, title: str) -> dict[str, Any]:
    header_fill = PatternFill("solid", fgColor=B9_COLOR_HEADER)
    key_fill = PatternFill("solid", fgColor=B9_COLOR_KPI)
    good_fill = PatternFill("solid", fgColor=B9_COLOR_GOOD)
    warn_fill = PatternFill("solid", fgColor=B9_COLOR_WARN)
    bad_fill = PatternFill("solid", fgColor=B9_COLOR_BAD)
    neutral_fill = PatternFill("solid", fgColor=B9_COLOR_NEUTRAL)
    thin = Side(style="thin", color=B9_COLOR_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet_kind = _visual_sheet_kind(title)
    tab_color = {
        "EJECUTIVA": B9_COLOR_HEADER,
        "BASE": "5B9BD5",
        "AUDITORIA": "A6A6A6",
        "MATRIZ": "70AD47",
        "DETALLE": "9EADCC",
    }.get(sheet_kind, B9_COLOR_HEADER)
    key_cols = {
        "SECTOR_ACADEMICO",
        "JORNADA",
        "NUEVO_ANTIGUO",
        "TASA_RESPUESTA_PCT",
        "PROMEDIO_AC_MEDIO",
        "NIVEL_AC_PROMEDIO",
    }
    pct_tokens = ("PCT", "PORC", "TASA")
    avg_tokens = ("PROMEDIO", "MEDIA", "MEDIO", "VALOR")
    count_tokens = ("ESTUDIANTES", "RESPONDIERON", "NO_RESPONDIERON", "CON_AC", "SIN_AC", "TOTAL", "FILAS", "COLUMNAS")

    ws.freeze_panes = "A2"
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        for cell in ws[1]:
            header = "" if cell.value is None else str(cell.value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = key_fill if header in key_cols else header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.sheet_properties.tabColor = tab_color
        try:
            ws.sheet_view.zoomScale = 90 if sheet_kind in {"EJECUTIVA", "AUDITORIA", "MATRIZ"} else 85
            ws.sheet_view.zoomScaleNormal = ws.sheet_view.zoomScale
            ws.sheet_view.view = "normal"
        except Exception:
            pass

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header = "" if ws.cell(1, col_idx).value is None else str(ws.cell(1, col_idx).value)
            header_upper = header.upper()
            max_len = len(header)
            width_cap = _visual_width_cap(header_upper, sheet_kind)
            for row_idx, cell in enumerate(ws[col_letter], start=1):
                text = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(text), 80))
                if row_idx == 1:
                    continue
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
                if any(token in header_upper for token in pct_tokens):
                    cell.number_format = "0.00"
                    if header_upper == "TASA_RESPUESTA_PCT" and isinstance(cell.value, (int, float)):
                        value = float(cell.value)
                        if value >= 60:
                            cell.fill = good_fill
                        elif value < 40:
                            cell.fill = bad_fill
                        else:
                            cell.fill = warn_fill
                elif any(token in header_upper for token in avg_tokens) or header_upper.startswith("ASI_") or header_upper.startswith("_AC_"):
                    cell.number_format = "0.00"
                elif any(token == header_upper or header_upper.endswith(token) for token in count_tokens):
                    cell.number_format = '#,##0'
                if row_idx % 2 == 0 and sheet_kind != "BASE" and cell.fill.fill_type is None:
                    cell.fill = neutral_fill
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, _visual_width_min(header_upper)), width_cap)

        ws.row_dimensions[1].height = 28

    return {
        "HOJA": title,
        "FREEZE_PANES": ws.freeze_panes,
        "AUTOFILTRO": ws.auto_filter.ref,
        "TABLA_EXCEL": "NO",
        "MOTIVO_TABLA_EXCEL": "Se replica BLOQUE 9: no crear ListObjects para evitar reparaciones en Excel Mac.",
        "COLOR_PESTANA": tab_color,
        "TIPO_HOJA": sheet_kind,
        "FILAS": ws.max_row,
        "COLUMNAS": ws.max_column,
        "FORMATO": "B9_REPLICADO_ENCABEZADO_AZUL_BORDES_FILTRO_FREEZE_ZOOM_ANCHOS_NUMERICOS",
    }


def _visual_sheet_kind(sheet_name: str) -> str:
    if sheet_name in EXECUTIVE_SHEETS:
        return "EJECUTIVA"
    if sheet_name in AUDIT_SHEETS or sheet_name.startswith("AUDITORIA"):
        return "AUDITORIA"
    if sheet_name == "BASE_SECTORIZADA":
        return "BASE"
    if sheet_name.startswith("MATRIZ"):
        return "MATRIZ"
    return "DETALLE"


def _visual_width_min(header_upper: str) -> int:
    if header_upper in {"SECTOR_ACADEMICO", "NUEVO_ANTIGUO", "TASA_RESPUESTA_PCT"}:
        return 18
    if header_upper in {"JORNADA", "RESPONDE"}:
        return 14
    return 10


def _visual_width_cap(header_upper: str, sheet_kind: str) -> int:
    if header_upper in {"CARRERA", "CARRERA_NORMALIZADA", "DETALLE", "RECOMENDACION", "VALOR", "CAMPO", "SEGMENTO"}:
        return 70
    if header_upper in {"SECTOR_ACADEMICO", "REGLA_SECTOR_ACADEMICO"}:
        return 38
    if sheet_kind == "BASE":
        return 42
    return 60


def apply_sectorized_format(path: Path, sheet_names: list[str]) -> list[dict[str, Any]]:
    wb = load_workbook(path)
    report: list[dict[str, Any]] = []
    for sheet_name in sheet_names:
        safe_sheet = sheet_name[:31]
        if safe_sheet not in wb.sheetnames:
            continue
        report.append(_format_sheet(wb[safe_sheet], safe_sheet))
    wb.save(path)
    wb.close()
    return report


def _auditoria_valor(df: pd.DataFrame, campo: str, default: str = "") -> str:
    if df.empty or not {"CAMPO", "VALOR"}.issubset(df.columns):
        return default
    sub = df[df["CAMPO"].astype(str).eq(campo)]
    if sub.empty:
        return default
    val = sub.iloc[0]["VALOR"]
    return "" if pd.isna(val) else str(val)


def _panel_df_block(ws, title: str, df: pd.DataFrame, start_row: int, start_col: int = 1) -> int:
    thin = Side(style="thin", color=B9_COLOR_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor=B9_COLOR_TITLE)
    header_fill = PatternFill("solid", fgColor=B9_COLOR_HEADER)
    max_cols = max(1, len(df.columns))
    end_col = start_col + max_cols - 1

    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=min(end_col, start_col + 8))
    cell_title = ws.cell(start_row, start_col, title)
    cell_title.fill = title_fill
    cell_title.font = Font(bold=True, size=12)
    cell_title.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    cell_title.border = border

    header_row = start_row + 1
    for offset, col in enumerate(df.columns, start=start_col):
        cell = ws.cell(header_row, offset, str(col))
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for r_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        row_values = list(row)
        format_hint = ""
        if "FORMATO" in df.columns:
            format_hint = str(row_values[list(df.columns).index("FORMATO")]).upper()
        for c_idx, value in enumerate(row, start=start_col):
            if pd.isna(value):
                value = ""
            cell = ws.cell(r_idx, c_idx, value)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=False)
            header = str(ws.cell(header_row, c_idx).value or "").upper()
            if header == "VALOR" and format_hint == "PORCENTAJE":
                cell.number_format = "0.00%"
            elif "TASA" in header or "PORC" in header or "PCT" in header:
                cell.number_format = "0.00"
            elif "PROMEDIO" in header or "NIVEL" in header:
                cell.number_format = "0.00"
            elif header in {"ESTUDIANTES", "RESPONDIERON", "NO_RESPONDIERON", "CON_AC", "SIN_AC", "VALOR"}:
                if isinstance(value, (int, float, np.integer, np.floating)):
                    cell.number_format = "#,##0"

    return header_row + len(df) + 2


def _panel_autofit(ws, max_width: int = 70) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        best = 0
        for cell in ws[letter]:
            text = "" if cell.value is None else str(cell.value)
            best = max(best, min(len(text), 90))
        ws.column_dimensions[letter].width = min(max(best + 2, 12), max_width)


def build_panel_insumos(path: Path) -> dict[str, pd.DataFrame]:
    base = pd.read_excel(path, sheet_name="BASE_SECTORIZADA")
    sector_jna = pd.read_excel(path, sheet_name="SECTOR_JORNADA_NA")
    carrera_jna = pd.read_excel(path, sheet_name="CARRERA_JORNADA_NA")
    lectura = pd.read_excel(path, sheet_name="LECTURA_EJECUTIVA_SECTOR")
    aud_cruce = pd.read_excel(path, sheet_name="AUDITORIA_CRUCE_AC")
    metadata: dict[str, Any] = {
        "anio": int(pd.to_numeric(base["ANIO_ANALISIS"], errors="coerce").dropna().iloc[0]),
        "periodo": int(pd.to_numeric(base["PERIODO_ANALISIS"], errors="coerce").dropna().iloc[0]),
        "fecha_ejecucion": str(base["FECHA_EJECUCION"].dropna().iloc[0]),
        "total_vigentes": int(len(base)),
    }
    for key in [
        "tokens_resultados_unicos",
        "tokens_duplicados_resultados_brutos",
        "filas_asociadas_duplicados",
        "tokens_ambiguos_con_resultado_no_atribuido",
        "tokens_resultados_no_asociados",
        "conflictos_ac_pendientes",
        "casos_con_fallback_llave",
    ]:
        value = _auditoria_valor(aud_cruce, key, "0")
        metadata[key] = int(float(value)) if str(value).strip() else 0
    kpis = build_panel_kpis(base, metadata)

    contexto = pd.DataFrame(
        [
            {"CAMPO": "Archivo base encuesta", "VALOR": _auditoria_valor(aud_cruce, "archivo_encuesta_usado", "NO_DISPONIBLE")},
            {"CAMPO": "Universo maestro", "VALOR": _auditoria_valor(aud_cruce, "fuente_universo", "DF_VIGENTES")},
            {"CAMPO": "Fecha/hora de generación", "VALOR": metadata["fecha_ejecucion"]},
            {"CAMPO": "Fuente de sectorización", "VALOR": _auditoria_valor(aud_cruce, "diccionario_sectorial", "OFICIAL")},
            {
                "CAMPO": "Nota",
                "VALOR": "Capa analítica adicional; no modifica el pipeline productivo original",
            },
        ]
    )
    sector = construir_resumen_sectorizado(base, ["SECTOR_ACADEMICO"], "SECTOR")

    lectura_panel = lectura.head(18).copy() if not lectura.empty else pd.DataFrame()
    criticos = carrera_jna[pd.to_numeric(carrera_jna["ESTUDIANTES"], errors="coerce").fillna(0) >= 10].copy()
    criticos = criticos.sort_values(["TASA_RESPUESTA_PCT", "ESTUDIANTES"], ascending=[True, False]).head(15)
    criticos = criticos[
        [
            "SECTOR_ACADEMICO",
            "CARRERA",
            "JORNADA",
            "NUEVO_ANTIGUO",
            "ESTUDIANTES",
            "RESPONDIERON",
            "TASA_RESPUESTA_PCT",
            "PROMEDIO_AC_MEDIO",
            "NIVEL_AC_PROMEDIO",
        ]
    ]
    matriz = sector_jna[
        [
            "SECTOR_ACADEMICO",
            "JORNADA",
            "NUEVO_ANTIGUO",
            "ESTUDIANTES",
            "RESPONDIERON",
            "TASA_RESPUESTA_PCT",
            "CON_AC",
            "PROMEDIO_AC_MEDIO",
            "NIVEL_AC_PROMEDIO",
        ]
    ].copy()
    trazabilidad = pd.DataFrame(
        [
            {"REFERENCIA": "Detalle individual", "HOJA": "BASE_SECTORIZADA"},
            {"REFERENCIA": "Trazabilidad carrera-sector", "HOJA": "AUDITORIA_SECTORES"},
            {"REFERENCIA": "Cardinalidad", "HOJA": "AUDITORIA_UNIVERSO_SECTORIZADO"},
            {"REFERENCIA": "Duplicados encuesta", "HOJA": "AUDITORIA_DUPLICADOS_ENCUESTA"},
            {"REFERENCIA": "Conflictos AC", "HOJA": "AUDITORIA_CONFLICTOS_AC"},
            {"REFERENCIA": "Reconciliación panel", "HOJA": "AUDITORIA_RECONCILIACION_PANEL"},
        ]
    )
    alertas = pd.DataFrame(
        [
            {"ALERTA": "Las hojas agregadas de encuesta se clasifican como AGREGADO_GLOBAL y no se atribuyen individualmente."},
            {"ALERTA": "La trayectoria se utiliza solo como fuente auxiliar AC; el universo proviene de DF_VIGENTES."},
            {"ALERTA": "Las preguntas clave no se sectorizan individualmente si no hay base individual disponible."},
            {"ALERTA": "Los segmentos pequeños deben interpretarse con cautela."},
            {"ALERTA": "Esta hoja resume; los datos completos siguen en las hojas fuente."},
        ]
    )
    return {
        "contexto": contexto,
        "kpis": kpis,
        "sector": sector,
        "lectura": lectura_panel,
        "criticos": criticos,
        "matriz": matriz,
        "trazabilidad": trazabilidad,
        "alertas": alertas,
    }


def insertar_panel_sectorizado_ejecutivo(workbook_path: Path) -> Path:
    if not workbook_path.exists():
        raise SystemExit(f"{PY_PREFIX}[ERROR] No existe libro para panel: {workbook_path}")

    insumos = build_panel_insumos(workbook_path)
    wb = load_workbook(workbook_path)
    if PANEL_SHEET in wb.sheetnames:
        del wb[PANEL_SHEET]
    ws = wb.create_sheet(PANEL_SHEET, 0)
    ws.sheet_properties.tabColor = "7030A0"
    ws.freeze_panes = "A4"
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.zoomScaleNormal = 90

    ws.merge_cells("A1:L1")
    metadata = {
        "anio": int(insumos["kpis"].loc[insumos["kpis"]["KPI"].eq("Año de análisis"), "VALOR"].iloc[0]),
        "periodo": int(insumos["kpis"].loc[insumos["kpis"]["KPI"].eq("Período de análisis"), "VALOR"].iloc[0]),
    }
    ws["A1"] = f"Panel Ejecutivo Sectorizado — Encuesta de Caracterización {metadata['anio']} P{metadata['periodo']}"
    ws["A1"].fill = PatternFill("solid", fgColor=B9_COLOR_HEADER)
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    ws["A2"] = "Lectura ejecutiva sectorizada basada en trayectoria AC y encuesta de caracterización."
    ws["A2"].font = Font(bold=True)
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    row_left = 4
    row_left = _panel_df_block(ws, "Contexto y fuente", insumos["contexto"], row_left, 1)
    row_left = _panel_df_block(ws, "Resumen por sector", insumos["sector"], row_left, 1)
    row_left = _panel_df_block(ws, "Lectura ejecutiva automática", insumos["lectura"], row_left, 1)
    row_left = _panel_df_block(ws, "Ranking de focos críticos", insumos["criticos"], row_left, 1)
    row_left = _panel_df_block(ws, "Matriz sector × jornada × nuevo/antiguo", insumos["matriz"], row_left, 1)
    row_left = _panel_df_block(ws, "Trazabilidad", insumos["trazabilidad"], row_left, 1)
    _panel_df_block(ws, "Alertas metodológicas", insumos["alertas"], row_left, 1)
    _panel_df_block(ws, "KPIs superiores", insumos["kpis"], 4, 10)

    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    _panel_autofit(ws)
    wb.save(workbook_path)
    wb.close()
    return workbook_path


def crear_panel_sectorizado_ejecutivo(input_path: Path, output_path: Path, abrir: bool = False) -> Path:
    if not input_path.exists():
        raise SystemExit(f"{PY_PREFIX}[ERROR] No existe archivo input: {input_path}")
    if output_path.resolve() == input_path.resolve():
        raise SystemExit(f"{PY_PREFIX}[ERROR] La salida no puede sobrescribir el archivo input.")

    shutil.copy2(input_path, output_path)
    insertar_panel_sectorizado_ejecutivo(output_path)
    if abrir:
        subprocess.run(["open", str(output_path)], check=False)
    return output_path


def write_output(encuesta: Path, salida: Path, sheets: dict[str, pd.DataFrame]) -> list[dict[str, Any]]:
    shutil.copy2(encuesta, salida)
    with pd.ExcelWriter(salida, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        for sheet_name in ALL_OUTPUT_SHEETS:
            df = sheets.get(sheet_name, pd.DataFrame())
            safe_sheet = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_sheet, index=False)
    return apply_sectorized_format(salida, ALL_OUTPUT_SHEETS)


def quality_audit_path(encuesta: Path, salida: Path) -> Path:
    anio, periodo, fecha = parse_period_from_name(encuesta)
    if anio and periodo and fecha:
        return salida.with_name(f"AUDITORIA_CALIDAD_SECTORIZADO_AC_{anio}_P{periodo}_{fecha}.xlsx")
    return salida.with_name(f"AUDITORIA_CALIDAD_{salida.stem}.xlsx")


def write_quality_audit(
    encuesta: Path,
    trayectoria: Path | None,
    salida: Path,
    auditoria: Path | None = None,
    visual_report: list[dict[str, Any]] | None = None,
) -> Path:
    """Auditoría dinámica: solo compara la salida con sus fuentes e invariantes."""
    auditoria = auditoria or quality_audit_path(encuesta, salida)
    xls = pd.ExcelFile(salida)
    output_sheet_names = list(xls.sheet_names)
    xls.close()
    base = pd.read_excel(salida, sheet_name="BASE_SECTORIZADA")
    aud_cruce = pd.read_excel(salida, sheet_name="AUDITORIA_CRUCE_AC")
    aud_universo = pd.read_excel(salida, sheet_name="AUDITORIA_UNIVERSO_SECTORIZADO")
    aud_duplicados = pd.read_excel(salida, sheet_name="AUDITORIA_DUPLICADOS_ENCUESTA")
    aud_descartes = pd.read_excel(salida, sheet_name="AUDITORIA_DESCARTES_ENCUESTA")
    aud_conflictos = pd.read_excel(salida, sheet_name="AUDITORIA_CONFLICTOS_AC")
    aud_panel = pd.read_excel(salida, sheet_name="AUDITORIA_RECONCILIACION_PANEL")

    original_xls = pd.ExcelFile(encuesta)
    original_sheets = list(original_xls.sheet_names)
    original_xls.close()
    missing_new = [sheet for sheet in NEW_SHEETS if sheet[:31] not in output_sheet_names]
    missing_granularity = [sheet for sheet in GRANULARITY_SHEETS if sheet[:31] not in output_sheet_names]
    missing_original = [sheet for sheet in original_sheets if sheet not in output_sheet_names]
    tiene_panel = PANEL_SHEET in output_sheet_names
    empty_rows = []
    for sheet in output_sheet_names:
        df_head = pd.read_excel(salida, sheet_name=sheet, nrows=5)
        empty_rows.append({"HOJA": sheet, "VACIA": df_head.empty, "COLUMNAS": len(df_head.columns)})
    empty = pd.DataFrame(empty_rows)
    total = int(len(base))
    official_total = int(float(_auditoria_valor(aud_cruce, "total_vigentes", str(total))))
    tokens_unique = int(float(_auditoria_valor(aud_cruce, "tokens_resultados_unicos", "0")))
    raw_result_rows = int(float(_auditoria_valor(aud_cruce, "filas_resultados_brutos", str(tokens_unique))))
    unique_result_rows = int(float(_auditoria_valor(aud_cruce, "filas_resultados_unicos", str(tokens_unique))))
    discarded_result_rows = int(float(_auditoria_valor(aud_cruce, "filas_descartadas_totales", str(len(aud_descartes)))))
    si = int(base["ESTADO_RESPUESTA"].eq("SI").sum())
    incomplete = int(base["ESTADO_RESPUESTA"].eq("INCOMPLETO").sum())
    no = int(base["ESTADO_RESPUESTA"].eq("NO").sum())
    present_mask = base["ESTADO_RESPUESTA"].isin(["SI", "INCOMPLETO"])
    associated_tokens = base.loc[present_mask, "_TOKEN_CRUCE_RESULTADO"].fillna("")
    associated_tokens_are_unique = bool(
        associated_tokens.ne("").all() and associated_tokens.is_unique
    )
    con_ac = int(base["CON_DATOS_AC"].fillna(False).sum())
    sin_ac = total - con_ac
    columnas_ac = pd.DataFrame(
        [{"COLUMNA": col, "PRESENTE": col in base.columns, "NO_NULOS": int(base[col].notna().sum()) if col in base.columns else 0} for col in EXPECTED_AC_COLS]
    )
    formato_visual = pd.DataFrame(visual_report or [])
    integridad = pd.DataFrame(
        [
            {"CHECK": "hojas_originales_preservadas", "OK": not missing_original, "DETALLE": ", ".join(map(str, missing_original))},
            {"CHECK": "hojas_nuevas_presentes", "OK": not missing_new, "DETALLE": ", ".join(missing_new)},
            {"CHECK": "hojas_granularidad_presentes", "OK": not missing_granularity, "DETALLE": ", ".join(missing_granularity)},
            {"CHECK": "panel_ejecutivo_primera_hoja", "OK": (not tiene_panel) or output_sheet_names[0] == PANEL_SHEET, "DETALLE": output_sheet_names[0]},
            {"CHECK": "universo_oficial_conservado", "OK": total == official_total, "DETALLE": f"base={total}; oficial={official_total}"},
            {"CHECK": "llave_final_unica", "OK": bool(base["LLAVE_ESTUDIANTE"].is_unique), "DETALLE": int(base["LLAVE_ESTUDIANTE"].duplicated().sum())},
            {"CHECK": "estados_reconcilian", "OK": si + incomplete + no == total, "DETALLE": f"SI={si}; INCOMPLETO={incomplete}; NO={no}; TOTAL={total}"},
            {"CHECK": "ac_reconcilia", "OK": con_ac + sin_ac == total, "DETALLE": f"CON={con_ac}; SIN={sin_ac}; TOTAL={total}"},
            {"CHECK": "respondientes_no_superan_tokens", "OK": si + incomplete <= tokens_unique, "DETALLE": f"respondientes={si + incomplete}; tokens={tokens_unique}"},
            {
                "CHECK": "procedencia_deduplicacion_reconcilia",
                "OK": raw_result_rows == unique_result_rows + discarded_result_rows == unique_result_rows + len(aud_descartes),
                "DETALLE": (
                    f"brutas={raw_result_rows}; canónicas={unique_result_rows}; "
                    f"descartadas_métrica={discarded_result_rows}; descartes_auditados={len(aud_descartes)}"
                ),
            },
            {"CHECK": "respuesta_no_multiplicada", "OK": associated_tokens_are_unique, "DETALLE": f"filas_asociadas={len(associated_tokens)}; tokens_asociados={associated_tokens.nunique()}"},
            {"CHECK": "sin_many_to_many", "OK": not bool(aud_universo["cardinalidad_observada"].astype(str).eq("many_to_many").any()), "DETALLE": ""},
            {"CHECK": "auditoria_universo_ok", "OK": bool(aud_universo["resultado"].astype(str).eq("OK").all()), "DETALLE": ""},
            {"CHECK": "reconciliacion_panel_ok", "OK": bool(aud_panel["Resultado"].astype(str).eq("OK").all()), "DETALLE": ""},
        ]
    )

    resumen = pd.DataFrame(
        [
            {"CAMPO": "archivo_auditado", "VALOR": str(salida)},
            {"CAMPO": "archivo_encuesta_base", "VALOR": str(encuesta)},
            {"CAMPO": "archivo_fuente_ac", "VALOR": str(trayectoria or "")},
            {"CAMPO": "fecha_hora_auditoria", "VALOR": _auditoria_valor(aud_cruce, "fecha_ejecucion", _auditoria_valor(aud_cruce, "fecha_hora_generacion", ""))},
            {"CAMPO": "resultado_general", "VALOR": "OK" if bool(integridad["OK"].all()) else "ERROR"},
        ]
    )

    with pd.ExcelWriter(auditoria, engine="openpyxl") as writer:
        pd.concat([resumen, pd.DataFrame([{"CAMPO": "resultado_integridad", "VALOR": "OK" if bool(integridad["OK"].all()) else "REVISAR"}])], ignore_index=True).to_excel(
            writer, sheet_name="VALIDACION_GENERAL", index=False
        )
        empty.to_excel(writer, sheet_name="HOJAS", index=False)
        columnas_ac.to_excel(writer, sheet_name="COLUMNAS_AC", index=False)
        formato_visual.to_excel(writer, sheet_name="FORMATO_VISUAL", index=False)
        integridad.to_excel(writer, sheet_name="DETALLE_CHECKS", index=False)
        aud_cruce.to_excel(writer, sheet_name="AUDITORIA_CRUCE_AC", index=False)
        aud_universo.to_excel(writer, sheet_name="AUDITORIA_UNIVERSO", index=False)
        aud_duplicados.to_excel(writer, sheet_name="DUPLICADOS_ENCUESTA", index=False)
        aud_descartes.to_excel(writer, sheet_name="DESCARTES_ENCUESTA", index=False)
        aud_conflictos.to_excel(writer, sheet_name="CONFLICTOS_AC", index=False)
        aud_panel.to_excel(writer, sheet_name="RECONCILIACION_PANEL", index=False)

    wb = load_workbook(auditoria)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_idx in range(1, ws.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = 0
                for cell in ws[col_letter]:
                    text = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(text))
                ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 70)
    wb.save(auditoria)
    wb.close()
    if not bool(integridad["OK"].all()):
        raise SectorizadoValidationError(
            "La auditoría dinámica del libro sectorizado contiene controles en ERROR."
        )
    return auditoria


def read_dataframe_path(
    path: Path,
    *,
    sheet_name: str | None = None,
    detect_structured_header: bool = False,
) -> pd.DataFrame:
    if not path.exists():
        raise SectorizadoValidationError(f"No existe la fuente tabular: {path}")
    suffix = path.suffix.lower()
    if suffix in {".pkl", ".pickle"}:
        data = pd.read_pickle(path)
        if not isinstance(data, pd.DataFrame):
            raise SectorizadoValidationError(f"El pickle no contiene un DataFrame: {path}")
        return data
    if suffix == ".csv":
        return pd.read_csv(path, sep=None, engine="python")
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        selected_sheet: str | int = sheet_name if sheet_name is not None else 0
        if sheet_name is not None:
            with pd.ExcelFile(path) as workbook:
                sheet_names = list(workbook.sheet_names)
            if sheet_name not in sheet_names:
                raise SectorizadoValidationError(f"No existe la hoja {sheet_name!r} en {path.name}.")
        if detect_structured_header:
            preview = pd.read_excel(path, sheet_name=selected_sheet, header=None, nrows=30)
            header_idx: int | None = None
            for idx, row in preview.iterrows():
                headers = {norm_header(value) for value in row.dropna().tolist()}
                if "CODCLI" in headers or ({"RUT", "TOKEN"} & headers):
                    header_idx = int(str(idx))
                    break
            if header_idx is None:
                raise SectorizadoValidationError(f"No se detectó cabecera institucional en {path.name}.")
            return pd.read_excel(path, sheet_name=selected_sheet, header=header_idx).dropna(how="all")
        return pd.read_excel(path, sheet_name=selected_sheet)
    raise SectorizadoValidationError(f"Formato tabular no soportado: {path.suffix}")


def cargar_resultados_canonicos_con_procedencia(
    resultados_path: Path,
    provenance_path: Path,
    *,
    token_col: str | None,
    submitdate_col: str | None,
    datestamp_col: str | None,
    response_id_col: str | None,
    key_kind: str,
) -> DedupResult:
    """Carga el contrato serializado del deduplicador institucional.

    El libro separa deliberadamente el frame canónico de la procedencia de
    sus filas físicas. Un consumidor sectorizado puede así reutilizar la misma
    selección sin volver a interpretar ni reducir los resultados.
    """
    if resultados_path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise SectorizadoValidationError(
            "--resultados-canonicos requiere un libro Excel con hoja RESULTADOS_CANONICOS."
        )
    if provenance_path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise SectorizadoValidationError(
            "--dedup-provenance debe ser un libro Excel con la procedencia canónica."
        )
    with pd.ExcelFile(resultados_path) as result_workbook:
        result_sheets = set(result_workbook.sheet_names)
    with pd.ExcelFile(provenance_path) as provenance_workbook:
        provenance_sheets = set(provenance_workbook.sheet_names)
    if "RESULTADOS_CANONICOS" not in result_sheets:
        raise SectorizadoValidationError(
            f"Falta RESULTADOS_CANONICOS en {resultados_path.name}."
        )
    required_provenance = {"METRICAS_DEDUP", "AUDITORIA_DEDUP", "DESCARTES_DEDUP"}
    missing = sorted(required_provenance - provenance_sheets)
    if missing:
        raise SectorizadoValidationError(
            f"Procedencia canónica incompleta en {provenance_path.name}: {', '.join(missing)}."
        )

    resultados = pd.read_excel(
        resultados_path,
        sheet_name="RESULTADOS_CANONICOS",
        dtype={"_TOKEN_RESULTADO": str},
    )
    metricas_df = read_dataframe_path(provenance_path, sheet_name="METRICAS_DEDUP")
    if not {"METRICA", "VALOR"}.issubset(metricas_df.columns):
        raise SectorizadoValidationError(
            "METRICAS_DEDUP debe contener las columnas METRICA y VALOR."
        )
    metric_names = metricas_df["METRICA"].fillna("").astype(str).str.strip()
    if metric_names.eq("").any() or metric_names.duplicated().any():
        raise SectorizadoValidationError(
            "METRICAS_DEDUP contiene nombres vacíos o duplicados."
        )
    try:
        metricas = {
            name: int(float(value))
            for name, value in zip(metric_names, metricas_df["VALOR"])
        }
    except (TypeError, ValueError) as exc:
        raise SectorizadoValidationError(
            "METRICAS_DEDUP contiene un valor no entero."
        ) from exc
    auditoria = pd.read_excel(
        provenance_path,
        sheet_name="AUDITORIA_DEDUP",
        dtype={
            "token": str,
            "ID seleccionado": str,
            "IDs descartados": str,
        },
    )
    descartes = pd.read_excel(
        provenance_path,
        sheet_name="DESCARTES_DEDUP",
        dtype={
            "token": str,
            "ID seleccionado": str,
            "ID descartado": str,
        },
    )
    return validar_resultados_canonicos_precomputados(
        resultados,
        metricas=metricas,
        auditoria=auditoria,
        descartes=descartes,
        token_col=token_col or "_TOKEN_RESULTADO",
        submitdate_col=submitdate_col,
        datestamp_col=datestamp_col,
        response_id_col=response_id_col,
        key_kind=key_kind,
    )


def write_failure_audit(
    requested_path: Path,
    error: Exception,
    *,
    anio: int | None,
    periodo: int | None,
    fecha_ejecucion: str | None,
) -> Path:
    requested_path.parent.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    failure_path = requested_path.with_name(f"{requested_path.stem}_FALLO_{stamp}.xlsx")
    traceback_completo = "".join(
        traceback.format_exception(type(error), error, error.__traceback__)
    )
    summary = pd.DataFrame(
        [
            {"CAMPO": "estado", "VALOR": "NO_GENERADO"},
            {"CAMPO": "bloque_fallido", "VALOR": "SECTORIZADO_AC"},
            {"CAMPO": "codigo_salida", "VALOR": 2},
            {"CAMPO": "año", "VALOR": anio},
            {"CAMPO": "período", "VALOR": periodo},
            {"CAMPO": "fecha_ejecución", "VALOR": fecha_ejecucion or ""},
            {"CAMPO": "tipo_error", "VALOR": type(error).__name__},
            {"CAMPO": "mensaje", "VALOR": str(error)},
            {"CAMPO": "traceback", "VALOR": traceback_completo},
            {"CAMPO": "salida_previa", "VALOR": "NO_SOBRESCRITA"},
            {"CAMPO": "salida_generada", "VALOR": "NO"},
        ]
    )
    traceback_df = pd.DataFrame(
        [
            {"linea": idx + 1, "texto": line}
            for idx, line in enumerate(traceback_completo.splitlines())
        ]
    )
    audit = error.audit if isinstance(error, SectorizadoValidationError) else pd.DataFrame()
    with pd.ExcelWriter(failure_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="FALLO_SECTORIZADO", index=False)
        traceback_df.to_excel(writer, sheet_name="TRACEBACK", index=False)
        if audit.empty:
            pd.DataFrame(columns=AUDIT_UNIVERSE_COLUMNS).to_excel(
                writer, sheet_name="AUDITORIA_UNIVERSO", index=False
            )
        else:
            audit.to_excel(writer, sheet_name="AUDITORIA_UNIVERSO", index=False)
    return failure_path


def generar_libro_sectorizado_atomico(
    *,
    encuesta: Path,
    trayectoria: Path | None,
    salida: Path,
    auditoria: Path,
    outputs: dict[str, pd.DataFrame],
    panel: bool,
) -> tuple[list[dict[str, Any]], Path]:
    """Publica ambos libros solo después de validar y restaura salidas previas ante fallo."""
    salida.parent.mkdir(parents=True, exist_ok=True)
    auditoria.parent.mkdir(parents=True, exist_ok=True)
    if salida.resolve() == auditoria.resolve():
        raise SectorizadoValidationError("La salida sectorizada y su auditoría deben usar rutas distintas.")
    with (
        tempfile.TemporaryDirectory(
            prefix="sectorizado_ac_publicacion_", dir=str(salida.parent)
        ) as output_tmp_dir,
        tempfile.TemporaryDirectory(
            prefix="sectorizado_ac_auditoria_", dir=str(auditoria.parent)
        ) as audit_tmp_dir,
    ):
        output_tmp_root = Path(output_tmp_dir)
        audit_tmp_root = Path(audit_tmp_dir)
        tmp_output = output_tmp_root / salida.name
        tmp_audit = audit_tmp_root / auditoria.name
        visual_report = write_output(encuesta, tmp_output, outputs)
        if panel:
            insertar_panel_sectorizado_ejecutivo(tmp_output)
            visual_report.append(
                {
                    "HOJA": PANEL_SHEET,
                    "FREEZE_PANES": "A4",
                    "AUTOFILTRO": "DINAMICO",
                    "TABLA_EXCEL": "NO",
                    "MOTIVO_TABLA_EXCEL": "Compatibilidad con Excel Mac.",
                    "COLOR_PESTANA": "7030A0",
                    "TIPO_HOJA": "EJECUTIVA",
                    "FILAS": "DINAMICO",
                    "COLUMNAS": "DINAMICO",
                    "FORMATO": "PANEL_EJECUTIVO_DINAMICO",
                }
            )
        write_quality_audit(
            encuesta,
            trayectoria,
            tmp_output,
            auditoria=tmp_audit,
            visual_report=visual_report,
        )
        previous_output = output_tmp_root / "salida_previa.xlsx"
        previous_audit = audit_tmp_root / "auditoria_previa.xlsx"
        output_existed = salida.exists()
        audit_existed = auditoria.exists()
        if output_existed:
            shutil.copy2(salida, previous_output)
        if audit_existed:
            shutil.copy2(auditoria, previous_audit)
        try:
            os.replace(tmp_output, salida)
            os.replace(tmp_audit, auditoria)
        except Exception:
            if output_existed:
                os.replace(previous_output, salida)
            else:
                salida.unlink(missing_ok=True)
            if audit_existed:
                os.replace(previous_audit, auditoria)
            else:
                auditoria.unlink(missing_ok=True)
            raise
    return visual_report, auditoria


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Crea análisis sectorizado desde el universo oficial vigente.")
    parser.add_argument("--input", type=Path, default=None, help="Libro sectorizado existente para agregar panel ejecutivo sin regenerar hojas")
    parser.add_argument("--output", type=Path, default=None, help="Salida para modo --input")
    parser.add_argument("--encuesta", "--excel-encuesta", dest="encuesta", type=Path, default=None, help="Ruta a ENCUESTA_CARACTERIZACION_*.xlsx")
    parser.add_argument("--universo-vigente", type=Path, default=None, help="Snapshot explícito de DF_VIGENTES")
    parser.add_argument("--resultados", type=Path, default=None, help="Resultados LimeSurvey individualizables")
    parser.add_argument(
        "--resultados-canonicos",
        action="store_true",
        help="Consume RESULTADOS_CANONICOS y su procedencia sin volver a deduplicar",
    )
    parser.add_argument(
        "--dedup-provenance",
        type=Path,
        default=None,
        help="Libro con METRICAS_DEDUP, AUDITORIA_DEDUP y DESCARTES_DEDUP",
    )
    parser.add_argument("--inscritos", type=Path, default=None, help="Inscritos LimeSurvey para auditoría")
    parser.add_argument(
        "--trayectoria",
        "--excel-trayectoria",
        dest="trayectoria",
        type=Path,
        default=None,
        help="Ruta a PREGRADO_TRAYECTORIA_CARACTERIZACION_*.xlsx",
    )
    parser.add_argument("--anio", type=int, default=None, help="Año parametrizado de la ejecución")
    parser.add_argument("--periodo", type=int, default=None, help="Período parametrizado de la ejecución")
    parser.add_argument("--fecha-ejecucion", type=str, default=None, help="Fecha/hora ISO de inicio de ejecución")
    parser.add_argument("--clave-resultados", type=str, default=None)
    parser.add_argument("--clave-inscritos", type=str, default=None)
    parser.add_argument("--tipo-llave-encuesta", type=str, default="token")
    parser.add_argument("--submitdate-col", type=str, default=None)
    parser.add_argument("--datestamp-col", type=str, default=None)
    parser.add_argument("--id-respuesta-col", type=str, default=None)
    parser.add_argument("--columnas-ac-requeridas", nargs="*", default=None)
    parser.add_argument("--minimo-variables-ac", type=int, default=1)
    parser.add_argument("--salida", dest="salida", type=Path, default=None, help="Ruta de salida .xlsx")
    parser.add_argument(
        "--auditoria",
        "--auditoria-output",
        dest="auditoria",
        type=Path,
        default=None,
        help="Ruta de auditoria de calidad .xlsx",
    )
    parser.add_argument("--panel", action="store_true", help="Agregar PANEL_SECTORIZADO_EJECUTIVO al libro generado")
    parser.add_argument("--abrir", action="store_true", help="Abrir archivo resultante al finalizar")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.input:
        input_path = args.input.expanduser()
        output_path = args.output.expanduser() if args.output else input_path.with_name(f"{input_path.stem}_PANEL.xlsx")
        result = crear_panel_sectorizado_ejecutivo(input_path, output_path, abrir=args.abrir)
        log(f"Archivo con panel ejecutivo generado: {result}")
        log("[VALIDACION] Panel dinámico generado=OK")
        return 0
    required = {
        "--encuesta": args.encuesta,
        "--universo-vigente": args.universo_vigente,
        "--resultados": args.resultados,
        "--anio": args.anio,
        "--periodo": args.periodo,
        "--fecha-ejecucion": args.fecha_ejecucion,
    }
    missing = [flag for flag, value in required.items() if value is None]
    if missing:
        raise SystemExit(f"{PY_PREFIX}[ERROR] Faltan parámetros obligatorios: {', '.join(missing)}")
    encuesta = args.encuesta.expanduser()
    universe_path = args.universo_vigente.expanduser()
    results_path = args.resultados.expanduser()
    inscritos_path = args.inscritos.expanduser() if args.inscritos else None
    trayectoria = args.trayectoria.expanduser() if args.trayectoria else None
    salida_arg = args.salida or args.output
    salida = salida_arg.expanduser() if salida_arg else default_output_path(encuesta)
    try:
        fecha_tag = pd.Timestamp(args.fecha_ejecucion).strftime("%Y%m%d")
    except Exception as exc:
        raise SystemExit(f"{PY_PREFIX}[ERROR] --fecha-ejecucion inválida: {args.fecha_ejecucion}") from exc
    auditoria = (
        args.auditoria.expanduser()
        if args.auditoria
        else salida.with_name(f"AUDITORIA_CALIDAD_SECTORIZADO_AC_{args.anio}_P{args.periodo}_{fecha_tag}.xlsx")
    )
    if salida.resolve() == encuesta.resolve():
        raise SystemExit(f"{PY_PREFIX}[ERROR] La salida no puede sobrescribir el archivo de encuesta original.")

    try:
        universo = read_dataframe_path(universe_path, detect_structured_header=universe_path.suffix.lower() not in {".pkl", ".pickle"})
        if args.dedup_provenance is not None and not args.resultados_canonicos:
            raise SectorizadoValidationError(
                "--dedup-provenance solo es válido junto con --resultados-canonicos."
            )
        dedup_precomputado: DedupResult | None = None
        if args.resultados_canonicos:
            provenance_path = (
                args.dedup_provenance.expanduser()
                if args.dedup_provenance is not None
                else results_path
            )
            dedup_precomputado = cargar_resultados_canonicos_con_procedencia(
                results_path,
                provenance_path,
                token_col=args.clave_resultados,
                submitdate_col=args.submitdate_col,
                datestamp_col=args.datestamp_col,
                response_id_col=args.id_respuesta_col,
                key_kind=args.tipo_llave_encuesta,
            )
            resultados = dedup_precomputado.resultados_unicos
        else:
            resultados = read_dataframe_path(results_path)
        inscritos = read_dataframe_path(inscritos_path) if inscritos_path is not None else None
        fuente_ac = (
            read_dataframe_path(trayectoria, sheet_name="TRAYECTORIA_POR_PERIODO")
            if trayectoria is not None
            else None
        )
        log(f"[PARAMETROS] Año={args.anio} | Periodo={args.periodo} | Fecha={args.fecha_ejecucion}")
        log(f"[UNIVERSO] Estudiantes vigentes oficiales={len(universo)}")
        sectorized = construir_base_sectorizada_oficial(
            universo,
            resultados,
            anio=args.anio,
            periodo=args.periodo,
            fecha_ejecucion=args.fecha_ejecucion,
            fuente_ac=fuente_ac,
            inscritos=inscritos,
            clave_resultados=args.clave_resultados,
            clave_inscritos=args.clave_inscritos,
            tipo_llave_encuesta=args.tipo_llave_encuesta,
            submitdate_col=args.submitdate_col,
            datestamp_col=args.datestamp_col,
            response_id_col=args.id_respuesta_col,
            columnas_ac_requeridas=args.columnas_ac_requeridas,
            minimo_variables_ac=args.minimo_variables_ac,
            execution_id=args.fecha_ejecucion,
            dedup_precomputado=dedup_precomputado,
        )
        outputs = build_outputs(encuesta, sectorized)
        log(
            f"[ENCUESTA] Filas brutas={sectorized.metadata['filas_resultados_brutos']} | "
            f"tokens únicos={sectorized.metadata['tokens_resultados_unicos']}"
        )
        log(f"[ENCUESTA] Tabla deduplicada={sectorized.metadata['filas_resultados_unicos']}")
        log(
            "[ENCUESTA] Tokens ambiguos con resultado no atribuido="
            f"{sectorized.metadata['tokens_ambiguos_con_resultado_no_atribuido']} | "
            "filas vigentes involucradas="
            f"{sectorized.metadata['filas_ambiguas_con_resultado_no_atribuido']}"
        )
        log(
            f"[RESPUESTA] SI={sectorized.metadata['total_si']} | "
            f"INCOMPLETO={sectorized.metadata['total_incompleto']} | "
            f"NO={sectorized.metadata['total_no']} | TOTAL={len(sectorized.base)}"
        )
        log(
            f"[AC] Con datos={sectorized.metadata['total_con_datos_ac']} | "
            f"Sin datos={sectorized.metadata['total_sin_datos_ac']} | TOTAL={len(sectorized.base)}"
        )
        log(f"[CARDINALIDAD] Base inicial={len(universo)} | Base final={len(sectorized.base)}")
        generar_libro_sectorizado_atomico(
            encuesta=encuesta,
            trayectoria=trayectoria,
            salida=salida,
            auditoria=auditoria,
            outputs=outputs,
            panel=args.panel,
        )
        log("[VALIDACION] Conservación de filas=OK")
        log("[VALIDACION] Reconciliación de estados=OK")
        log("[VALIDACION] Reconciliación AC=OK")
        log("[VALIDACION] Sin cruces many-to-many=OK")
        log(f"Archivo sectorizado generado: {salida}")
        log(f"Archivo auditoría generado: {auditoria}")
        if args.abrir:
            subprocess.run(["open", str(salida)], check=False)
            subprocess.run(["open", str(auditoria)], check=False)
        return 0
    except Exception as exc:
        failure = write_failure_audit(
            auditoria,
            exc,
            anio=args.anio,
            periodo=args.periodo,
            fecha_ejecucion=args.fecha_ejecucion,
        )
        log(f"[ERROR] Archivo sectorizado no generado: {exc}")
        log(f"[ERROR] Auditoría del fallo: {failure}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
