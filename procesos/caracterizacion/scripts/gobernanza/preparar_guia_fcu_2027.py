"""Extract source questions and build an explicit, reviewable FCU crosswalk.

This prepares JSON only. Workbook authoring is performed by the JS builder.
"""
from pathlib import Path
from collections import Counter
import hashlib
import json
import re
import runpy
import sys
import unicodedata
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'gobernanza/FCU/2027/preparacion'
MANUALS = {2026: ROOT/'docs/Manual_de_Aplicacion_FCU_2026.txt',
           2027: BASE/'evidencias/Manual_de_Aplicacion_FCU_2027.txt'}
INVENTORY = BASE/'evidencias/Resumen_clasificacion_preguntas_FCU_2026.xlsx'

def norm(text):
    return ''.join(c for c in unicodedata.normalize('NFD', text.lower())
                   if c.isalnum() and not unicodedata.combining(c))

def choices(block):
    labels=[]
    for line in block.splitlines():
        if line.startswith(('꙱','⃝')):
            label=line.lstrip('꙱⃝ ').split('[T:')[0].strip()
            if label and any(c.isalpha() for c in label):
                labels.append(re.sub(r'\s+Nota \d+.*$', '',label))
    return labels

RESOURCE_ITEMS = ['Teléfono inteligente o smartphone','Notebook o computador portátil','Tablet',
 'Computador de escritorio (PC)','Internet Wifi','Internet con plan de datos móviles en tu teléfono móvil',
 'Cámara web (integrada a un dispositivo o externa)','Micrófono (integrado a un dispositivo o externo)',
 'Parlantes y/o audífonos','Impresora']
MINEDUC_ITEMS = ['La información sobre gratuidad, becas y beneficios fue clara.',
 'El calendario de postulación y matrícula era fácil de encontrar.',
 'El buscador de carreras de acceso.mineduc.cl fue útil','El buscador de carreras de mifuturo.cl fue útil',
 'El proceso de llenado del FUAS (Formulario Único de Acreditación Socioeconómica) fue fácil de hacer.']

def raw_manual(year):
    lines = MANUALS[year].read_text().splitlines()
    start = lines.index('5.3 FCU en Plataformas Propias')
    end = lines.index(f'5.4 Cambios en la FCU {year}', start)
    pages = {}
    page = None
    for i, line in enumerate(lines):
        match = re.match(r'--- Página (\d+) ---', line)
        if match:
            page = int(match[1])
        pages[i] = page
    anchors = []
    seen = set()
    for i in range(start+1, end):
        line = lines[i]
        match = re.match(r'^(\d{1,2}\.\d{1,2})\.?\s+(.+)$', line)
        if not match:
            match = re.match(r'^(\d{1,2})\.\s+(.+)$', line)
        if match:
            ref, text = match.groups()
            if text.startswith('Pregunta ') or ref in seen:
                continue
            seen.add(ref)
            anchors.append((i, 'P'+ref, text))
        elif line.startswith('Señala tu grado de acuerdo'):
            block = re.search(r'\((\d)/\d\)', line)
            if block:
                parent = 'P45' if year == 2026 else 'P39'
                anchors.append((i, parent+'.B'+block[1], line))
    rows = []
    index = {}
    for pos, (i, ref, first) in enumerate(anchors):
        stop = anchors[pos+1][0] if pos+1 < len(anchors) else end
        text_lines = [re.sub(r'^\[CONDICIONALIDAD[^]]*\]\s*', '', first)]
        for j in range(i+1, stop):
            line = lines[j].strip()
            if line.startswith('[') or line.startswith(('Selecciona', 'Marca ', 'Acepto ')) or line.startswith(('꙱', '⃝', '(1)')):
                break
            if re.match(r'^--- Página|^\d+$|^C$|^Nota \d+', line) or not line:
                continue
            if re.match(r'^[IVX]+\.', line):
                break
            text_lines.append(line)
        text = ' '.join(text_lines).split('[Pregunta')[0].split('[ Pregunta')[0].strip()
        row = [ref, 'Bloque sin numeración propia' if '.B' in ref else ('Pregunta condicional/componente' if '.' in ref else 'Pregunta principal'), text, pages[i], i+1, MANUALS[year].name]
        rows.append(row)
        index[ref] = {'text': text, 'page': pages[i], 'line': i+1,
                      'block': '\n'.join(lines[i:stop]), 'row': len(rows)+1}
    # Consent statements are source rows, not guessed LimeSurvey IDs.
    for i in range(start, end):
        if lines[i].startswith('Acepto que estos datos sean compartidos') or lines[i].startswith(('꙱ Autorizo a la institución', '꙱ He leído y acepto')):
            text = lines[i].lstrip('꙱ ')
            j = i+1
            while j < end and not lines[j].startswith(('⃝', '꙱', '---', 'CONSENTIMIENTO')) and not re.match(r'^\d',lines[j]) and lines[j].strip():
                text += ' '+lines[j]
                j += 1
            ref = 'CONS'+str(sum(r[1]=='Consentimiento' for r in rows)+1)
            rows.append([ref,'Consentimiento',text,pages[i],i+1,MANUALS[year].name])
    assert {r[0] for r in rows if r[1]=='Pregunta principal'} == {f'P{i}' for i in range(1,46 if year==2026 else 40)}
    return rows, index

# Explicit version correspondence; numbering is never used as a fuzzy match.
VERSION = {'P1':None,'P2':None,'P3':None,'P4':'P1','P4.1':'P1.1','P5':'P2','P6':'P3','P6.1':'P3.1','P7':None,
 'P8':'P4','P9':'P5','P9.1':None,'P10':'P6','P10.1':None,'P10.2':None,'P11':'P7','P11.1':'P7.1','P11.2':'P7.2','P11.3':None,
 'P12':'P8','P13':'P9','P14':'P10','P15':'P11','P16':'P12','P17':'P14','P17.1':'P14.1','P17.2':None,'P18':'P15','P19':None,'P20':'P13',
 'P21':'P16','P22':'P17','P23':'P18','P24':'P19','P25':'P20','P26':'P21','P27':'P22','P28':'P23','P29':'P24','P30':'P25','P31':'P26',
 'P32':'P27','P33':'P28','P34':'P29','P35':'P30','P36':'P31','P37':'P32','P37.1':'P32.1','P38':'P33','P39':'P34','P40':'P35','P40.1':'P35.1','P40.2':None,
 'P41':'P36','P42':'P36','P43':'P37','P44':'P38','P45':'P39','P45.B2':'P39.B2','P45.B3':'P39.B2'}
VERSION.update({f'P12.{i}':f'P8.{i}' for i in range(1,7)})
VERSION.update({f'P25.{i}':f'P20.{i}' for i in range(1,6)})
VERSION.update({f'P27.{i}':f'P22.{i}' for i in range(1,5)})

# These mappings identify content, not an assertion about live conditions.
OFFICIAL = {
 'G01Q07':'P6','G02Q02':'P6.1','G00Q02':'P8','G01Q03':'P9','G07Q04':'P9.1','G01Q74':'P9.1',
 'G01Q75':'P10','G01Q05':'P10','G01Q06':'P10.1','G01Q14':'P10.2','G01Q82':'P11','G01Q17':'P11.1','G01Q30':'P11.2','G01Q31':'P11.3',
 'G01Q33':'P12.6','G02Q76':'P15','G02Q08':'P15','G01Q10':'P17','G01Q90':'P17.1','G00Q11':'P17.2',
 'G00Q12':'P18','G02Q15':'P18','G02Q17':'P19','G01Q18':'P20','G01Q19':'P21','G02Q95':'P22',
 'G03Q28':'P23','G03Q30':'P24','G03Q31':'P25','G03Q33':'P25.1','G03Q44':'P25.2','G01Q45':'P25.3','G03Q46':'P25.4','G03Q47':'P25.5',
 'G03Q48':'P26','G03Q51':'P27','G03Q54':'P27.1','G03Q55':'P27.2','G03Q56':'P27.3','G03Q57':'P27.4',
 'G04Q37':'P30','G04Q39':'P31','G01Q50':'P32','G01Q41':'P33','G01Q43':'P34','G01Q55':'P35','G01Q84':'P36',
 'G06Q61':'P37','G06Q97':'P37.1','G01Q62':'P38','G01Q63':'P39','G01Q64':'P40','G01Q65':'P40.1','G06Q66':'P40.2',
 'G01Q67':'P41','G01Q68':'P42','G01Q69':'P43','G01Q70':'P44','G07Q71':'P45','G01Q72':'P45.B2','G01Q73':'P45.B3'}
INSTITUTIONAL = {'G01Q09','G01Q13','G01Q24','G01Q25','G01Q26','G01Q35','G01Q36','G01Q58','G01Q59','G01Q87',
 'G02Q20','G02Q21','G02Q22','G02Q23','G02Q27','G02Q74','G02Q75','G02Q89','G03Q76','G05Q57','G05Q60','G05Q86'}
CONSENTS = {'G00Q91','G00Q92','G00Q93','G00Q94'}
CHANGED = {'P8','P13','P20','P33','P35','P36','P41','P42','P43','P45','P45.B2','P45.B3'}
SPECIAL = {
 'P8': 'Cambiar alternativas a Hombre, Mujer y No binario.',
 'P13': 'Validar 11 dígitos: 56 + 9 + ocho dígitos; ejemplo 56912345678.',
 'P20': 'Reemplazar categorías de jefatura por la pregunta directa y respuestas Sí/No.',
 'P33': 'Fusionar ferias/charlas y prensa física/digital; reordenar las 13 opciones oficiales. Conservar las opciones institucionales adicionales fuera del bloque oficial.',
 'P35': 'Mantener cuatro ítems; retirar la categoría neutral de la escala, usar 1–4.',
 'P36': 'Mantener cinco ítems; retirar neutral y conservar 0 = No lo utilicé más 1–4 de acuerdo.',
 'P41': 'Cambiar acceso binario por frecuencia 1 Nunca, 2 A veces, 3 Frecuentemente, 4 Siempre.',
 'P42': 'Integrar espacio de estudio como ítem 11 de P36; retirar el formato independiente solo después de crear ese ítem.',
 'P43': 'Hacer obligatoria la pregunta directa de selección múltiple; incorporar No cuento con red de apoyo y revisar exclusividad.',
 'P45': 'Reestructurar como dos bloques y 15 afirmaciones, con escala 1–4. Consultar las filas de cada SQ antes de eliminar ítems.',
 'P45.B2': 'Reubicar/reformular afirmaciones en los dos bloques oficiales; consultar las filas de cada SQ.',
 'P45.B3': 'Reubicar/reformular afirmaciones en los dos bloques oficiales; consultar las filas de cada SQ.'}

NEW_DISPOSITION = [
 'Me esfuerzo por terminar todo lo que empiezo','He superado momentos difíciles para alcanzar una meta importante',
 'Cuando estoy en un entorno educacional o laboral, me siento parte de una comunidad',
 'Soy capaz de sacrificar el descanso y la diversión para dedicarme a mis estudios.',
 'Pongo mucho esfuerzo en mis estudios y actividades académicas', 'Realizo mis actividades académicas con dedicación y a conciencia.',
 'Estoy seguro(a) de que si estudio y me organizo, puedo obtener las notas que quiero','Me siento seguro(a) de la carrera que elegí',
 'Tengo dificultades para mantener el foco en proyectos que duran más de unos meses',
 'Cuando me va mal en algo, me cuesta manejar emociones fuertes como la rabia, la pena o el nerviosismo.',
 'Dudo que con mis habilidades académicas tenga un buen desempeño en la carrera.',
 'Cuando tengo un desafío académico, me cuesta organizarme para resolverlo.',
 'Tengo dificultades para comprender los textos que leo','Se me hace difícil expresar mis ideas por escrito.',
 'No me siento tan confiado(a) en que podré titularme']
OLD_DISPOSITION = [
 'Me gusta participar de actividades extraprogramáticas, como talleres, voluntariados u otros.',
 'Me gusta leer y tengo el hábito de hacerlo.', 'Tengo dificultad para comprender un texto escrito o expresarme a través de la escritura.',
 'Termino todo lo que empiezo', 'Tengo dificultades para mantener el foco en proyectos que duran más de unos meses.',
 'He superado momentos difíciles para conquistar una meta importante.',
 'No me siento cómodo(a) conociendo personas nuevas.', 'Cuando tengo un conflicto con otras personas, busco considerar todos los puntos de vista.',
 'Cuando estoy en entornos educacionales o laborales, tiendo a sentirme parte de una comunidad',
 'Cuando me va mal en algo, me cuesta manejar emociones fuertes como la rabia o la pena',
 'Soy capaz de sacrificar la diversión para dedicarme a mis estudios','Pongo mucho esfuerzo en mis estudios y actividades académicas.',
 'Dudo que con mis habilidades académicas tenga un buen desempeño en mi carrera.',
 'Cuando estudio o tengo actividades académicas, lo hago a conciencia y dedicadamente.',
 'Estoy seguro/a de que si estudio y me organizo puedo obtener las notas que deseo.',
 'Cuando no entendía algo o no me iba bien en el colegio, buscaba ayuda para que me explicaran lo que no sabía',
 'No me siento tan confiado/a en que terminaré mis estudios y me titularé',
 'Cuando tengo que enfrentar un desafío académico, me cuesta organizarme y resolver los problemas que aparecen.',
 'Me siento seguro/a de la decisión que tomé al elegir mi carrera e institución']
DIS_TARGET = {'G07Q71': [[],[],[13,14],[1],[9],[2]],
              'G01Q72':[[],[],[3],[10],[4],[5]],
              'G01Q73':[[11],[6],[7],[],[15],[12],[8]]}

def main(destination):
    raws, idx = {}, {}
    for year in [2026,2027]:
        raws[year], idx[year] = raw_manual(year)
    # Literal matrix items are source enunciados, not new LimeSurvey keys.
    item_sets={(2026,'P35'):['Me pareció fácil','Me pareció claro','Me pareció rápido','Me dieron una buena atención'],
               (2027,'P30'):['Me pareció fácil','Me pareció claro','Me pareció rápido','Me dieron una buena atención'],
               (2026,'P36'):MINEDUC_ITEMS,(2027,'P31'):MINEDUC_ITEMS,
               (2026,'P41'):RESOURCE_ITEMS,
               (2027,'P36'):RESOURCE_ITEMS+['Un ambiente y espacio físico que permite la concentración y que sea cómodo para estudiar']}
    for (year,ref),labels in item_sets.items():
        for i,label in enumerate(labels,1):
            assert norm(label) in norm(idx[year][ref]['block']),(year,ref,label)
            raws[year].append([f'{ref}.I{i}','Ítem de matriz',label,idx[year][ref]['page'],None,MANUALS[year].name])
    # Read the source-only inventory; never promote its historical categories silently.
    w = load_workbook(INVENTORY, read_only=True, data_only=True)
    inventory = {r[0]: list(r) for r in w['Codigo_x_Manual'].iter_rows(min_row=5,values_only=True) if r[0]}
    assert len(inventory)==103
    inst = []
    for rowno, row in enumerate(w['Institucionales'].iter_rows(min_row=2,values_only=True),2):
        if row[3]=='PREGUNTA_INSTITUCIONAL_ADICIONAL':
            inst.append([row[1],row[2],row[3], 'Resumen_clasificacion_preguntas_FCU_2026.xlsx / Institucionales', rowno])
    assert len(inst)==46
    options = {}
    for rowno,row in enumerate(w['Opciones 2026'].iter_rows(min_row=6,values_only=True),6):
        if row[0] in inventory and row[1] and row[6] and 'Alternativa declarada' in row[6]:
            options[(row[0],row[1])] = (row[3],rowno)
    # Visible location is evidence; unseen group order is left unconfirmed.
    namespace = runpy.run_path(str(ROOT/'scripts/gobernanza/ordenar_matriz_fcu_2027.py'), run_name='navigation_metadata')
    locations = namespace['POSITIONS']
    xml_path=BASE/'evidencias/LimeSurvey_141234_grupo_643_20261006.lsg'
    if not xml_path.exists():
        xml_path=Path('/Users/alexi/Downloads/limesurvey_group_643.lsg')
    live={}
    if xml_path.exists():
        xml=ET.parse(xml_path).getroot()
        for row in xml.findall('./questions/rows/row'):
            q={c.tag:c.text for c in row}
            if q.get('parent_qid')=='0':live[q['title']]=q
    assert live.get('G01Q74',{}).get('relevance','').find('G07Q04')>=0, 'Need evidence for the name-social detail control'
    cross=[]
    def add(code, sub, text, oldref=None, newref=None, category='VERIFICAR ORIGEN', method='', change='', final='', source=''):
        inv=inventory.get(code,[])
        loc=locations.get(code,(None,'UBICACIÓN PENDIENTE',None,'No visible en capturas'))
        old=idx[2026].get(oldref,{})
        new=idx[2027].get(newref,{})
        if not final:final=new.get('text','') or ('Fuera del bloque oficial FCU 2027.' if category=='ELIMINAR FCU' else text)
        correction=''
        if inv and inv[5]=='PREGUNTA_INSTITUCIONAL_ADICIONAL' and oldref:
            correction='Inventario anterior contradicho por pregunta literal del Manual 2026; no usar esa clasificación heredada.'
        evidence=f'Manual 2026 {oldref}, página {old.get("page")}, línea {old.get("line")}; Manual 2027 {newref}, página {new.get("page")}, línea {new.get("line")}' if oldref and newref else source
        if oldref and not newref:evidence=f'Manual 2026 {oldref}, página {old.get("page")}; retiro explícito Manual 2027 §5.4, páginas 41–43.'
        if category=='CONSERVAR INSTITUCIONAL':change='Mantener este contenido institucional. No eliminarlo por la actualización FCU.'
        if category=='VERIFICAR ORIGEN':change='No eliminar ni modificar: confirmar el vínculo del código con el Manual y su función en la estructura actual.'
        cross.append([f'C{len(cross)+1:03d}',code,sub,text,inv[5] if inv else '',oldref or '',old.get('text',''),newref or '',new.get('text',''),method,correction,category,change,final,loc[1],loc[2],evidence or source])
    for code, inv in inventory.items():
        text=inv[1]
        if code in OFFICIAL:
            oldref=OFFICIAL[code];newref=VERSION[oldref];assert oldref in idx[2026]
            if newref:assert newref in idx[2027]
            changed_choices = newref and [norm(c) for c in choices(idx[2026][oldref]['block'])] != [norm(c) for c in choices(idx[2027][newref]['block'])]
            category='ELIMINAR FCU' if not newref else ('MODIFICAR FCU' if oldref in CHANGED or changed_choices or norm(idx[2026][oldref]['text'])!=norm(idx[2027][newref]['text']) else 'CONSERVAR FCU')
            change= ('Retirar esta pregunta del bloque oficial FCU 2027.' if category=='ELIMINAR FCU' else SPECIAL.get(oldref,'Actualizar el enunciado al texto literal de 2027.' if category=='MODIFICAR FCU' else 'Mantener la pregunta oficial; actualizar su referencia de numeración.'))
            if code=='G01Q74':change='Retirar únicamente el campo de nombre social vinculado a G07Q04. El LSG confirma que se muestra cuando G07Q04 = AO01.'
            final = idx[2027][newref]['text'] if newref else ''
            if newref and category=='MODIFICAR FCU':
                if (2027,newref) in item_sets:
                    final += '\nÍtems oficiales: '+' | '.join(item_sets[(2027,newref)])
                elif choices(idx[2027][newref]['block']):
                    final += '\nOpciones oficiales (encabezados del Manual): '+' | '.join(choices(idx[2027][newref]['block']))
            add(code,'',text,oldref,newref,category,'Correspondencia explícita de contenido; códigos de LimeSurvey por inventario. '+('Vínculo al padre confirmado por LSG 643.' if code=='G01Q74' else ''),change,final,source='Comparación de manuales y hoja Codigo_x_Manual.')
        elif code in CONSENTS:
            add(code,'',text,category='CONSERVAR FCU',method='Texto oficial de consentimiento repetido en ambos manuales. Rama exacta debe verificarse.',change='Conservar el consentimiento; comprobar condición y alcance.',source='Manuales 2026 y 2027 §5.3; encabezados de consentimiento. LSG 643 confirma estos cuatro controles.')
        elif code in INSTITUTIONAL:
            add(code,'',text,category='CONSERVAR INSTITUCIONAL',method='Inventario institucional y ausencia de correspondencia numerada directa en ambos manuales.',source='RAW Institucionales: código exacto del inventario; no inferir retiro por ausencia en FCU.')
        else:
            add(code,'',text,method='Cruce anterior ambiguo o control desagregado/filtro sin equivalencia confirmada.',source='Codigo_x_Manual / clasificación 2026. No se usa un mejor candidato automático.')
    # Raw options for information sources: their literal source rows support the item-level join.
    choice_refs={}
    for year, ref in [(2026,'P33'),(2027,'P28')]:
        for number,line in enumerate(idx[year][ref]['block'].splitlines()):
            if line.startswith('꙱'):
                label=line.lstrip('꙱ ').strip();rid=f'{ref}.O{len([x for x in choice_refs if x[0]==year])+1}'
                raws[year].append([rid,'Opción de pregunta',label,idx[year][ref]['page'],idx[year][ref]['line']+number,MANUALS[year].name])
                choice_refs[(year,label)]=rid
    assert len([x for x in choice_refs if x[0]==2026])==15
    assert len([x for x in choice_refs if x[0]==2027])==13
    old_labels=[k[1] for k in choice_refs if k[0]==2026]
    new_labels=[k[1] for k in choice_refs if k[0]==2027]
    medium_map=[(0,3),(1,1),(2,2),(3,0),(4,4),(5,4),(6,5),(7,6),(8,7),(9,8),(10,9),(11,9),(12,10),(13,11),(14,12)]
    for old_i,new_i in medium_map:
        oldlabel,newlabel=old_labels[old_i],new_labels[new_i]
        key=f'SQ{old_i+2:03d}'
        if ('G01Q41',key) not in options:continue
        label,rowno=options[('G01Q41',key)]
        add('G01Q41',key,label,'P33','P28','MODIFICAR FCU','Opción observada en encabezado 2026, cotejada con listas literales de ambos manuales.',
            ('Fusionar con la otra opción de ferias/charlas.' if old_i in [4,5] else 'Fusionar prensa física y digital.' if old_i in [10,11] else 'Conservar el contenido y ubicarlo en el orden oficial 2027.'),
            f'Opción {new_i+1}: {newlabel}',f'Opciones 2026 fila {rowno}; RAW FCU 2026 {choice_refs[(2026,oldlabel)]}; RAW FCU 2027 {choice_refs[(2027,newlabel)]}')
        cross[-1][16]=f'RAW FCU 2026 {choice_refs[(2026,oldlabel)]}; RAW FCU 2027 {choice_refs[(2027,newlabel)]}; Opciones 2026 fila {rowno}'
        cross[-1][5:9]=[choice_refs[(2026,oldlabel)],oldlabel,choice_refs[(2027,newlabel)],newlabel]
    for key in [('G01Q41','SQ017'),('G01Q41','other')]:
        if key in options:
            label,rowno=options[key];add(key[0],key[1],label,category='CONSERVAR INSTITUCIONAL',method='Opción observada que no figura en las 15 opciones oficiales 2026 ni en las 13 de 2027.',
                final='Conservar en el contenido institucional; excluir de las 13 opciones del bloque oficial FCU.',source=f'Opciones 2026 fila {rowno}; listas oficiales P33 (2026) y P28 (2027).')
    # Disposition: every removal is paired to a literal former item, not a PES position.
    disposition_refs={}
    for year, values in [(2026,OLD_DISPOSITION),(2027,NEW_DISPOSITION)]:
        cleaned=norm('\n'.join(idx[year][r]['block'] for r in (['P45','P45.B2','P45.B3'] if year==2026 else ['P39','P39.B2'])))
        for i,label in enumerate(values,1):
            assert norm(label) in cleaned,(year,i,label)
            rid=f'DIS{year}.{i}';disposition_refs[(year,i)]=rid
            raws[year].append([rid,'Ítem de matriz',label,36 if year==2026 and i<=12 else 37 if year==2026 else 38 if i<=3 else 39,None,MANUALS[year].name])
    old_number=0
    for code, targets in DIS_TARGET.items():
        for k,tg in enumerate(targets,2):
            old_number+=1;key=f'SQ{k:03d}';label,rowno=options[(code,key)]
            category='MODIFICAR FCU' if tg else 'ELIMINAR FCU'
            target='\n'.join(f'{disposition_refs[(2027,n)]}: {NEW_DISPOSITION[n-1]}' for n in tg) if tg else 'El ítem no está en los 15 enunciados oficiales 2027.'
            add(code,key,label,OFFICIAL[code],VERSION[OFFICIAL[code]],category,'Comparación explícita de afirmación 2026 con los dos bloques 2027.',
                'Dividir el ítem en dos afirmaciones y usar escala 1–4.' if len(tg)==2 else 'Reformular/reubicar esta afirmación y usar escala 1–4.' if tg else 'Retirar SOLO esta afirmación del bloque oficial; conservar los ítems con equivalencia.',target)
            cross[-1][16]=f'RAW FCU 2026 {disposition_refs[(2026,old_number)]}; RAW FCU 2027 '+(', '.join(disposition_refs[(2027,n)] for n in tg) if tg else 'P39 y bloque 2: no hay equivalente')
            cross[-1][5:9]=[disposition_refs[(2026,old_number)],OLD_DISPOSITION[old_number-1],
                           ' | '.join(disposition_refs[(2027,n)] for n in tg),
                           '\n'.join(NEW_DISPOSITION[n-1] for n in tg)]
    handled={('G01Q41',k) for _,k in options if _=='G01Q41'}|{(code,f'SQ{k:03d}') for code,targets in DIS_TARGET.items() for k in range(2,len(targets)+2)}
    parents={r[1]:r for r in cross if not r[2]}
    for (code,key),(label,rowno) in options.items():
        if (code,key) in handled:continue
        parent=parents[code]
        oldref=OFFICIAL.get(code);newref=VERSION.get(oldref)
        category=parent[11]
        change=parent[12]
        final=label
        if oldref=='P41':final=f'{RESOURCE_ITEMS[int(key[2:])-2]}: 1 Nunca, 2 A veces, 3 Frecuentemente, 4 Siempre.'
        elif oldref=='P35':final=f'{item_sets[(2027,"P30")][int(key[2:])-2]}: 1 Muy en desacuerdo, 2 En desacuerdo, 3 De acuerdo, 4 Muy de acuerdo.'
        elif oldref=='P36':final=f'{MINEDUC_ITEMS[int(key[2:])-1]}: 0 No lo utilicé; 1–4 de acuerdo.'
        elif oldref=='P43':
            newopts=choices(idx[2027]['P37']['block'])
            if any(norm(label)==norm(x) for x in newopts):
                category='CONSERVAR FCU';change='Mantener esta alternativa: aparece literalmente en la lista 2027. No retirar por el cambio de nombres del mapa PES.'
            elif key=='SQ002':
                category='MODIFICAR FCU';change='Actualizar el texto de la primera alternativa al oficial 2027.';final=newopts[0]
            else:
                category='VERIFICAR ORIGEN';change='Confirmar correspondencia de esta opción con la nueva lista; no retirar por posición o nombre de campo PES.'
        elif newref and category not in ['ELIMINAR FCU','CONSERVAR INSTITUCIONAL']:
            target_labels=choices(idx[2027][newref]['block'])
            equivalent=next((x for x in target_labels if norm(x)==norm(label)),None)
            if equivalent:
                category='CONSERVAR FCU';change='Mantener esta opción: figura en la lista oficial 2027.';final=equivalent
            else:
                category='VERIFICAR ORIGEN';change='Confirmar la equivalencia de esta SQ con las alternativas/ítems oficiales. No heredar la acción del padre a una opción no conciliada.'
                final='Pendiente de correspondencia individual; conservar el elemento actual mientras se coteja.'
        add(code,key,label,oldref,newref,category,'Subpregunta/opción observada en encabezados; asociación con pregunta padre, no identidad de nuevo código.',change,final,source=f'Opciones 2026 fila {rowno}')
        cross[-1][16]+=f'; Opciones 2026 fila {rowno}'
        if oldref in ['P35','P36','P41']:
            item=int(key[2:])-(1 if oldref=='P36' else 2)
            cross[-1][5:9]=[f'{oldref}.I{item+1}',item_sets[(2026,oldref)][item],
                           f'{newref}.I{item+1}',item_sets[(2027,newref)][item]]
    # Missing instrument controls are explicitly not converted into guessed deletions.
    represented=set(OFFICIAL.values())
    for ref,source in idx[2026].items():
        if ref in represented:continue
        nr=VERSION[ref]
        add('SIN CÓDIGO CONFIRMADO','',source['text'],ref,nr,'VERIFICAR ORIGEN','Pregunta del Manual sin código identificado en el inventario disponible.',source='RAW FCU 2026 / RAW FCU 2027. Localizar el control o la fuente externa antes de editar.')
    cross.sort(key=lambda r: (locations.get(r[1],(99,'',999,''))[0] or 99,locations.get(r[1],(99,'',999,''))[2] or 999,r[1],r[2]))
    # Stable row identity allows guide-to-crosswalk references after sorting.
    for i,row in enumerate(cross,1):row[0]=f'C{i:03d}'
    guide=[]
    for i,r in enumerate(cross,2):
        instruction=r[11]+'\nQué hacer: '+r[12]+'\nCómo queda: '+r[13]
        guide.append([instruction,r[1],r[2],r[14],r[15],r[3],r[11],r[0],f'Cruces!A{i}:Q{i}'])
    categories=Counter(r[11] for r in cross)
    assert all(cat in {'ELIMINAR FCU','MODIFICAR FCU','CONSERVAR FCU','CONSERVAR INSTITUCIONAL','VERIFICAR ORIGEN'} for cat in categories)
    assert all(r[5] or r[1] in {'G01Q41'} for r in cross if r[11]=='ELIMINAR FCU')
    assert len({r[1] for r in cross if r[1]!='SIN CÓDIGO CONFIRMADO'})==103
    sheets=[
      {'name':'RAW FCU 2026','headers':['ID fuente (no código LimeSurvey)','Tipo','Pregunta literal / enunciado','Página del PDF fuente','Línea del TXT fuente','Fuente'],'rows':raws[2026]},
      {'name':'RAW FCU 2027','headers':['ID fuente (no código LimeSurvey)','Tipo','Pregunta literal / enunciado','Página del PDF fuente','Línea del TXT fuente','Fuente'],'rows':raws[2027]},
      {'name':'RAW Institucionales','headers':['Código LimeSurvey','Pregunta literal del inventario','Clasificación original (sin corregir)','Fuente','Fila fuente'],'rows':inst},
      {'name':'Cruces','headers':['ID cruce','Código LimeSurvey','Clave subpregunta/opción','Texto LimeSurvey fuente 2026','Clasificación inventario anterior','Referencia raw FCU 2026','Pregunta literal FCU 2026','Referencia raw FCU 2027','Pregunta literal FCU 2027','Método de correspondencia','Discrepancia del inventario','Categoría resultante','Cambio concreto','Cómo queda en 2027','Grupo LimeSurvey','Posición visible','Evidencia'],'rows':cross},
      {'name':'Guía para editar','headers':['Qué hacer / cambio concreto / cómo queda','Código pregunta','Clave subpregunta/opción','Grupo LimeSurvey','Posición visible','Elemento actual','Categoría','ID cruce','Referencia al cruce'],'rows':guide}]
    sources=[{'path':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [*MANUALS.values(),INVENTORY]]
    if xml_path.exists():sources.append({'path':str(xml_path),'sha256':hashlib.sha256(xml_path.read_bytes()).hexdigest()})
    payload={'sheets':sheets,'categories':dict(categories),'sources':sources,'rules':'Manual literal antes de decisiones; el raw institucional preserva la clasificación heredada, los errores se resuelven en Cruces. Los IDs DIS/P..O son identificadores documentales, no códigos LimeSurvey. La actualización de nombres/campos PES no autoriza eliminar opciones presentes en el Manual 2027.'}
    Path(destination).write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'sheets':{s['name']:len(s['rows']) for s in sheets},'categories':dict(categories)},ensure_ascii=False))

if __name__=='__main__':
    main(sys.argv[1])
