"""Integrate completed period analyses without changing their frozen rules or files."""
from collections import Counter, defaultdict
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from caracterizacion_limesurvey import normalize_cross_value
import enriquecer_analisis_anual as enrichment
import paneles_analisis_anual as dashboards
from completar_fuentes_excel_anual import complete_source_sheets, validate_source_dependencies

POLICY = 'ANUAL_INTEGRADO_2026-v2'
DIMENSIONS = ['Período', 'Sector', 'Modalidad', 'Sede', 'Carrera', 'Jornada']
EMPTY = {'', 'nan', 'none', 'null', 'n/a', 'na', '-', '--', '.', 's/i', 'sin info'}
YES = {'sí', 'si', 'yes', 'y', '1', 'true'}
NO = {'no', 'n', '0', 'false'}
BASE_HEADERS = ['Año', *DIMENSIONS, 'Persona', 'CODCLI', 'Estado respuesta',
                'ID respuesta', 'Promedio notas 2026', 'Nivel AC', 'Nuevo/antiguo en carrera',
                'Fuente modalidad', 'Archivo base', 'Fila fuente', 'Incluir (fx)',
                'Completa (fx)', 'Parcial (fx)', 'Sin respuesta (fx)',
                'Suma promedio (fx)', 'N promedio (fx)']


def text(value):
    if value is None or pd.isna(value):
        return ''
    return str(value).strip()


def identifier(value):
    value = text(value)
    return re.sub(r'\.0+$', '', value) if re.fullmatch(r'\d+\.0+', value) else value


def person_identity(row):
    token = normalize_cross_value(identifier(row['TOKEN_NORMALIZADO'])).lstrip('0')
    rut = identifier(row.get('RUT_NORMALIZADO')).replace('.', '').replace(' ', '').upper().lstrip('0')
    if not rut:
        return token
    if '-' in rut:
        return rut.split('-', 1)[0]
    if rut == token:
        return token
    # BASE_SECTORIZADA uses normalizar_rut_con_dv: compact RUT, retaining its DV.
    if re.fullmatch(r'\d+[0-9K]', rut):
        return rut[:-1]
    raise ValueError('RUT_NORMALIZADO incompatible con el contrato de la base sectorizada.')


def sha(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def col(index):
    output = ''
    while index >= 0:
        index, remainder = divmod(index, 26)
        output = chr(65 + remainder) + output
        index -= 1
    return output


def ref(sheet, column, start, end=None):
    name = sheet.replace("'", "''")
    return f"'{name}'!${col(column)}${start}" + (f':${col(column)}${end}' if end is not None else '')


def spec(name, headers, rows, note='', widths=None):
    return {'name': name, 'headers': headers, 'rows': rows, 'note': note,
            'widths': widths or [], 'formulas': {}, 'percent_columns': []}


def read_base(path):
    preview = pd.read_excel(path, sheet_name='BASE_SECTORIZADA', header=None, nrows=20, dtype=str)
    for i, row in preview.iterrows():
        if {'TOKEN_NORMALIZADO', 'ESTADO_RESPUESTA', 'PERIODO_ANALISIS'}.issubset(set(row.dropna())):
            data = pd.read_excel(path, sheet_name='BASE_SECTORIZADA', header=i, dtype=str).dropna(how='all')
            data['_FILA_ANUAL'] = data.index + int(i) + 2
            return data
    raise ValueError(f'{path.name}: no se encontró el encabezado de BASE_SECTORIZADA.')


def collect_period(entry, run_dir, resources=None, academic_cache=None):
    for role, source in entry['snapshot_sources'].items():
        if sha(source) != entry['source_hashes'][role]:
            raise ValueError(f'P{entry["periodo"]}: cambió la copia de entrada {role}.')
    candidates = sorted(Path(run_dir).glob('02_SALIDAS/*_ANALISIS_SECTORIZADO_AC.xlsx'))
    if len(candidates) != 1:
        raise ValueError(f'P{entry["periodo"]}: se esperaba una salida sectorizada, se encontraron {len(candidates)}.')
    path = candidates[0]
    base = read_base(path)
    results = pd.read_excel(entry['snapshot_sources']['RUTA_RESULTADOS'],dtype=str,keep_default_na=False)
    policy, catalog, inventory_path = resources or enrichment.load_resources()
    inventory = enrichment.shared_inventory(entry, results.columns, policy, inventory_path)
    grade_path = entry['snapshot_sources']['RUTA_PROMEDIOS']
    grade_hash = entry['source_hashes']['RUTA_PROMEDIOS']
    cache = {} if academic_cache is None else academic_cache
    if grade_hash not in cache:
        cache[grade_hash] = enrichment.academic_metrics(pd.read_excel(grade_path, sheet_name='Hoja1', dtype=str), policy)
    metrics, grade_quality, grade_periods = cache[grade_hash]
    base, enrichment_audit = enrichment.enrich_base(base, entry, catalog, policy, metrics)
    print(f'[ANUAL][P{entry["periodo"]}] Enriquecimiento: registros={len(base)}; '
          f'modalidad respaldada={sum(r[5] != "No determinada" for r in enrichment_audit)}; '
          f'sin sector={sum(r[4] == "Sin clasificar" for r in enrichment_audit)}; '
          f'con notas validas 2026={sum(r[11] > 0 for r in enrichment_audit)}. '
          'Universo y archivos originales conservados.', flush=True)
    return {'entry': entry, 'base': base, 'results': results, 'base_path': str(path),
            'inventory': inventory, 'inventory_path': str(inventory_path), 'base_hash': sha(path),
            'comparability': 'Comparable P1/P2/P3: instrumento comun 2026 autorizado; codigos/textos verificados',
            'enrichment_audit': enrichment_audit,
            'grade_quality': [[f'P{entry["periodo"]}', name, value] for name, value in grade_quality],
            'grade_periods': grade_periods,
            'canonical_labels': dict(json.loads((ROOT/policy['instrumento']['columnas_path']).read_text(encoding='utf-8'))),
            'governed_sources': [enrichment.POLICY_PATH, enrichment.verified_resource(policy['catalogo']),
                                 ROOT/policy['instrumento']['columnas_path']]}


def modality(row):
    # MODALIDAD in DATOSDEALUMNOS describes secondary school, not the IPSS program.
    for field in ('MODALIDAD_CARRERA', 'MODALIDAD_PROGRAMA'):
        if text(row.get(field)):
            return text(row[field]), text(row.get('FUENTE_MODALIDAD_ANUAL')) or field
    return 'No determinada', 'Sin campo de modalidad de carrera; no se usa MODALIDAD escolar ni se infiere desde jornada'


def question_column(header):
    match = re.match(r'^((?:G\d+Q\d+|Q\d+)(?:\[[^\]]+\])*)(?:\.\s*|\s+|$)', header)
    return (match.group(1), header[match.end():].strip()) if match else None


def prepare(periods):
    """The official period universe is the input; do not rerun eligibility or dedup rules."""
    if sorted(p['entry']['periodo'] for p in periods) != [1, 2, 3]:
        raise ValueError('La consolidación requiere exactamente P1, P2 y P3.')
    records, responses, dictionary, issues, sources = [], [], [], [], []
    for part in sorted(periods, key=lambda p: p['entry']['periodo']):
        entry, base, raw = part['entry'], part['base'], part['results']
        period = entry['periodo']
        if entry['anio'] != 2026 or entry['survey_id'] != {1:'923565',2:'443299',3:'967488'}[period]:
            raise ValueError('Perfil de año/encuesta distinto del anual 2026.')
        required = {'TOKEN_NORMALIZADO', 'SECTOR_ACADEMICO', 'ESTADO_RESPUESTA', 'CARRERA',
                    'JORNADA', 'SEDE', 'ANIO_ANALISIS', 'PERIODO_ANALISIS', 'ID_RESPUESTA_SELECCIONADO'}
        if missing := required - set(base):
            raise ValueError(f'P{period}: faltan columnas de base: {sorted(missing)}')
        id_columns = [c for c in raw if c.lower().split('.')[0].strip() == 'id']
        token_columns = [c for c in raw if c.lower().split('.')[0].strip() == 'token']
        if len(id_columns) != 1:
            raise ValueError(f'P{period}: resultados sin ID único identificable.')
        if len(token_columns) != 1:
            raise ValueError(f'P{period}: resultados sin token identificable.')
        ids = raw[id_columns[0]].map(identifier)
        if ids.eq('').any() or ids.duplicated().any():
            raise ValueError(f'P{period}: ID de respuesta vacío o repetido; no elegir una fila arbitraria.')
        id_rows = {v: i for i, v in enumerate(ids)}
        selected = defaultdict(list)
        for _, row in base.iterrows():
            if (identifier(row['ANIO_ANALISIS']), identifier(row['PERIODO_ANALISIS'])) != (str(entry['anio']), str(period)):
                raise ValueError('La base sectorizada contiene otro año/período.')
            identity = person_identity(row)
            if not identity:
                raise ValueError(f'P{period}: registro sin identidad, no se puede contar personas únicas.')
            state = text(row['ESTADO_RESPUESTA'])
            if state not in {'SI', 'INCOMPLETO', 'NO'}:
                raise ValueError(f'P{period}: estado de respuesta desconocido {state!r}.')
            rid = identifier(row['ID_RESPUESTA_SELECCIONADO'])
            if state != 'NO' and rid not in id_rows:
                raise ValueError(f'P{period}: respuesta seleccionada no encontrada en su export original.')
            if state != 'NO':
                raw_token = normalize_cross_value(identifier(raw.iloc[id_rows[rid]][token_columns[0]]))
                selected_token = normalize_cross_value(identifier(row.get('_TOKEN_CRUCE_RESULTADO',row['TOKEN_NORMALIZADO'])))
                if raw_token != selected_token:
                    raise ValueError(f'P{period}: ID seleccionado y token no corresponden a la misma respuesta.')
            mode, mode_source = modality(row)
            mean = pd.to_numeric(row.get('PROMEDIO_AC_REFERENCIA'), errors='coerce')
            records.append([entry['anio'], f'P{period}', text(row['SECTOR_ACADEMICO']) or 'Sin clasificar',
                            mode, text(row['SEDE']) or 'Sin dato', text(row['CARRERA']) or 'Sin dato',
                            text(row['JORNADA']) or 'Sin dato', identity, identifier(row.get('CODCLI')),
                            state, rid, None if pd.isna(mean) else float(mean),
                            text(row.get('TRAMO_NIVEL_AC')) or 'Sin dato', text(row.get('NUEVO_ANTIGUO')),
                            mode_source, Path(part['base_path']).name, int(row.get('_FILA_ANUAL', len(records) + 7))])
            if rid and state != 'NO':
                selected[rid].append(len(records) - 1)
            if text(row.get('TOKEN_AMBIGUO_UNIVERSO')).upper() in {'TRUE', '1', 'SI'}:
                issues.append([f'P{period}', identity, 'TOKEN_AMBIGUO', 'Se conserva la decisión del maestro; no se reatribuye la respuesta.'])
        for rid, indices in selected.items():
            if len({records[i][7] for i in indices}) != 1:
                raise ValueError(f'P{period}: una respuesta se atribuye a distintas personas.')
        for rid,i in id_rows.items():
            if rid not in selected:
                issues.append([f'P{period}',identifier(raw.iloc[i][token_columns[0]]),'RESPUESTA_NO_SELECCIONADA',
                               f'ID {rid}; fila {i+2} del export. No se atribuye al universo. Motivo detallado en auditorías originales del período.'])
        columns = []
        for j, header in enumerate(raw.columns):
            parsed = question_column(str(header))
            if not parsed:
                continue
            code, label = parsed
            label = label or part.get('canonical_labels', {}).get(code, '')
            parent = code.split('[')[0]
            q = part['inventory'].get(parent, {})
            kind = q.get('type', '')
            valid = [text(v) for v in raw.iloc[:, j] if text(v).lower() not in EMPTY]
            choices = list(dict.fromkeys(valid))
            is_open = kind in {'S', 'T', 'U', 'Q'} or '[other]' in code.lower() or (kind=='P' and 'comment' in code.lower())
            aggregate = not is_open and (bool(kind) or len(choices) <= 30)
            status = part.get('comparability', 'Separada por período; equivalencia no aprobada')
            dictionary.append([f'P{period}', entry['survey_id'], code, q.get('group', 'Sin grupo verificado'),
                               label or q.get('text', code), q.get('type_label', 'Tipo no verificado'),
                               status, str(header), part['inventory_path'] or 'Falta inventario propio de este período'])
            columns.append({'column': j, 'code': code, 'label': label or q.get('text', code),
                            'group': q.get('group', 'Sin grupo verificado'), 'kind': kind,
                            'parent': parent,
                            'aggregate': aggregate, 'choices': choices, 'open': is_open,
                            'status': status})
        if len({c['code'] for c in columns}) != len(columns):
            raise ValueError(f'P{period}: columnas de pregunta con código repetido; requiere conciliación.')
        if not columns and len(raw):
            raise ValueError(f'P{period}: no se identificaron columnas de preguntas en el export.')
        for question in columns:
            question['multiple'] = question['kind'] in {'M','P'} and not question['open']
            question['percent_verified'] = bool(question['kind'])
            if question['multiple']:
                siblings = [q['column'] for q in columns if q['parent']==question['parent'] and not q['open']]
                unexpected = [v for c in siblings for v in raw.iloc[:,c] if text(v).lower() not in EMPTY|YES|NO]
                if unexpected:
                    raise ValueError(f'P{period} {question["parent"]}: marcador múltiple no reconocido; revisar el inventario/export.')
                question['parent_valid'] = [any(text(row[c]).lower() in YES for c in siblings) for row in raw.values.tolist()]
        responses.append({'period': period, 'id_column': list(raw.columns).index(id_columns[0]), 'headers': list(raw.columns), 'rows': raw.astype(object).where(raw.notna(), None).values.tolist(),
                          'columns': columns, 'selected': {id_rows[rid]: indices for rid, indices in selected.items()}})
        for var, path in entry.get('snapshot_sources', {}).items():
            sources.append([f'P{period}', entry['survey_id'], var, str(path), entry.get('source_hashes', {}).get(var, '')])
        sources.append([f'P{period}', entry['survey_id'], 'BASE_SECTORIZADA', part['base_path'], part.get('base_hash','')])
        if part['inventory_path']:
            sources.append([f'P{period}',entry['survey_id'],'INVENTARIO',part['inventory_path'],sha(part['inventory_path'])])
        for path in part.get('governed_sources', []):
            sources.append([f'P{period}',entry['survey_id'],'REGLA_ANUAL',str(path),sha(path)])
    for identity, count in Counter(r[7] for r in records).items():
        if count > 1:
            issues.append(['Anual', identity, 'VARIAS_MATRICULAS', f'{count} registros; personas únicas por selección, sin eliminar matrículas.'])
    return {'year': 2026, 'policy': POLICY, 'records': records, 'responses': responses,
            'dictionary': dictionary, 'issues': issues, 'sources': sources,
            'enrichment_audit': [r for p in periods for r in p.get('enrichment_audit', [])],
            'grade_quality': [r for p in periods for r in p.get('grade_quality', [])],
            'grade_periods': sorted({v for p in periods for v in p.get('grade_periods', [])})}


def reference_counts(model, selection=None):
    selection = selection or {}
    rows = [r for r in model['records'] if all(r[i + 1] in values for i, values in selection.items())]
    people = {r[7] for r in rows}
    complete = {r[7] for r in rows if r[9] == 'SI'}
    partial = {r[7] for r in rows if r[9] == 'INCOMPLETO'} - complete
    return [len(rows), len(people), len(complete), len(partial), len(people - complete - partial)]


def criterion(cell):
    return f'"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({cell},"~","~~"),"*","~*"),"?","~?")'


def make_payload(model):
    records = model['records']
    filters = []
    for d, label in enumerate(DIMENSIONS):
        values = sorted({r[d + 1] for r in records})
        if d == 0:
            values = ['P1', 'P2', 'P3']
        for value in values:
            filters.append([label, value, 'Sí'])
    controls = spec('Filtros', ['Dimensión', 'Valor', 'Incluir'], filters,
                    'Una selección compartida por todas las vistas. Marque Sí/No por valor; admite cualquier combinación. Cero casos no amplía la selección.', [22, 65, 16])
    controls['validations'] = [{'range': f'C7:C{len(filters)+6}', 'values': ['Sí', 'No']}]
    # Immutable cohort profiles share one filter calculation; no repeated full-base scans.
    profile_keys = sorted({tuple(row[1:7]) for row in records})
    profile_index = {key:i+7 for i,key in enumerate(profile_keys)}
    profiles = spec('Perfiles selección', [*DIMENSIONS,'Incluir (fx)'],
                    [list(key)+[None] for key in profile_keys],
                    'Combinaciones observadas en el corte archivado. Solo editar Sí/No en Filtros; regenerar para cambiar las fuentes.',
                    [12,30,24,14,55,16,18])
    for key,r in profile_index.items():
        cells = [f'COUNTIFS({ref("Filtros",0,7,len(filters)+6)},"{label}",{ref("Filtros",1,7,len(filters)+6)},{criterion(f"{col(d)}{r}")},{ref("Filtros",2,7,len(filters)+6)},"Sí")=1'
                 for d,label in enumerate(DIMENSIONS)]
        profiles['formulas'][f'G{r}'] = '=IF(AND('+','.join(cells)+'),1,0)'

    def selected_profiles(indices, state=None):
        rows = sorted({profile_index[tuple(records[i][1:7])] for i in indices
                       if state is None or records[i][9] == state})
        return 'MAX('+','.join(ref(profiles['name'],6,r) for r in rows)+')' if rows else '0'
    base = spec('Base anual', BASE_HEADERS, [r + [None] * 6 for r in records],
                'Una fila por registro del universo oficial de cada período. Las matrículas no son personas únicas.', [12, 12, 30, 24, 14, 55, 16, 18])
    for i, row in enumerate(records, 7):
        base['formulas'].update({f'R{i}': '='+ref(profiles['name'],6,profile_index[tuple(row[1:7])]),
                                f'S{i}': f'=IF(J{i}="SI",R{i},0)',
                                f'T{i}': f'=IF(J{i}="INCOMPLETO",R{i},0)',
                                f'U{i}': f'=IF(J{i}="NO",R{i},0)',
                                f'V{i}': f'=IF(ISNUMBER(L{i}),L{i}*R{i},0)',
                                f'W{i}': f'=IF(ISNUMBER(L{i}),R{i},0)'})
    memberships = defaultdict(list)
    for i, row in enumerate(records):
        memberships[('Anual', 'Todos', row[7])].append(i)
        for d, label in enumerate(DIMENSIONS):
            memberships[(label, row[d + 1], row[7])].append(i)
        memberships[('Cruce', json.dumps(row[1:4], ensure_ascii=False), row[7])].append(i)
        memberships[('Sector período', json.dumps(row[1:3], ensure_ascii=False), row[7])].append(i)
        memberships[('Sector carrera', json.dumps([row[2],row[5]], ensure_ascii=False), row[7])].append(i)
    count_rows, count_formulas = [], {}
    for (dimension, value, person), indices in sorted(memberships.items()):
        r = len(count_rows) + 7
        count_rows.append([dimension, value, person, None, None, None, None])
        sums = [selected_profiles(indices,state) for state in (None,'SI','INCOMPLETO')]
        count_formulas.update({f'D{r}': f'=IF({sums[0]}>0,1,0)',
                               f'E{r}': f'=IF({sums[1]}>0,1,0)',
                               f'F{r}': f'=IF(E{r}=1,0,IF({sums[2]}>0,1,0))',
                               f'G{r}': f'=D{r}-E{r}-F{r}'})
    counts = spec('Conteos personas', ['Dimensión', 'Valor', 'Persona', 'Incluida (fx)', 'Completa (fx)', 'Parcial (fx)', 'Sin respuesta (fx)'], count_rows,
                  'Persona única dentro de cada segmento. Completa prevalece sobre parcial. Sectores/carreras pueden solaparse; no sumar sus personas.')
    counts['formulas'] = count_formulas
    last = max(7, len(count_rows) + 6)
    order = ['Anual', 'Período', 'Sector', 'Modalidad', 'Sede', 'Carrera', 'Jornada', 'Cruce', 'Sector período', 'Sector carrera']
    segments = sorted({(r[0], r[1]) for r in count_rows}, key=lambda item: (order.index(item[0]), item[1]))
    participation_rows = []
    for dimension, value in segments:
        dimensions = json.loads(value) if dimension == 'Cruce' else [value if dimension == d else '' for d in DIMENSIONS[:3]]
        participation_rows.append([dimension, ' / '.join(dimensions) if dimension == 'Cruce' else value,
                                   None, None, None, None, None, *dimensions])
    participation = spec('Participación', ['Dimensión', 'Segmento', 'Personas', 'Completas', 'Incompletas', 'Sin respuesta', '% completas', 'Período', 'Sector IPSS', 'Modalidad'],
                         participation_rows,
                         'Personas únicas por segmento y cruce Período/Sector IPSS/Modalidad. Filtros compartidos; segmentos solapados no se suman.', [24, 58, 18, 18, 18, 18, 18, 14, 32, 24])
    for r in range(7, len(segments) + 7):
        indices = [i+7 for i,row in enumerate(count_rows) if tuple(row[:2]) == segments[r-7]]
        for j in range(4):
            participation['formulas'][f'{col(j+2)}{r}'] = f'=SUM({ref("Conteos personas",j+3,indices[0],indices[-1])})'
        participation['formulas'][f'G{r}'] = f'=IF(C{r}=0,"No aplica",D{r}/C{r})'
    participation['percent_columns'] = [6]
    summary = spec('Resumen', ['Indicador', 'Selección actual'], [[label, None] for label in
                   ['Registros de matrícula', 'Personas únicas', 'Personas con respuesta completa', 'Personas con respuesta incompleta', 'Personas sin respuesta', 'Participación completa']],
                   '2026 P1/P2/P3. No es la carga FCU cerrada. Datos personales: uso interno. Selección compartida en Filtros.', [52, 25])
    summary['formulas']['B7'] = f'=SUM({ref("Base anual",17,7,max(7,len(records)+6))})'
    for offset in range(4):
        end = max(7,len(segments)+6)
        summary['formulas'][f'B{8+offset}'] = f'=SUMIFS({ref("Participación",offset+2,7,end)},{ref("Participación",0,7,end)},"Anual",{ref("Participación",1,7,end)},"Todos")'
    summary['formulas']['B12'] = '=IF(B8=0,"No aplica",B9/B8)'
    summary['percent_cells'] = ['B12']
    summary['chart'] = {'range': 'A9:B11', 'start': 'D7', 'end': 'J21', 'title': 'Estado de respuesta de personas únicas'}
    grade_periods = ', '.join(model.get('grade_periods', [])) or 'sin cobertura validada'
    performance = spec('Rendimiento', ['Período ingreso', 'Sector', 'Nivel AC', 'Matrículas con notas 2026', 'Media de promedios 2026'], [],
                       f'Notas finales 2026 disponibles: períodos académicos {grade_periods}. Cada matrícula promedia sus notas válidas; aquí se promedian matrículas, no personas ni semestres. No es cierre anual. Promedio previo en auditoría.', [18, 34, 24, 26, 26])
    for period, sector, level in sorted({(r[1], r[2], r[12]) for r in records}):
        r = len(performance['rows']) + 7
        performance['rows'].append([period,sector,level,None,None])
        end = max(7,len(records)+6)
        criteria = f'{ref("Base anual",1,7,end)},A{r},{ref("Base anual",2,7,end)},{criterion(f"B{r}")},{ref("Base anual",12,7,end)},{criterion(f"C{r}")}'
        performance['formulas'][f'D{r}'] = f'=SUMIFS({ref("Base anual",22,7,end)},{criteria})'
        performance['formulas'][f'E{r}'] = f'=IF(D{r}=0,"Sin notas disponibles al corte",SUMIFS({ref("Base anual",21,7,end)},{criteria})/D{r})'
    question_headers = ['Período', 'Grupo', 'Código literal', 'Pregunta / subpregunta', 'Valor observado', 'N respuestas', 'Base válida', '% de respuesta / selección', 'Sin respuesta válida a esta pregunta', 'Comparabilidad entre períodos']
    questions = spec('Preguntas', question_headers, [],
                     'Solo completas del universo. Múltiple: base con alguna selección en la pregunta principal. Sin respuesta válida incluye vacíos/saltos; no distingue omisión de no aplica. Tipo desconocido: sin porcentaje.', [12,30,20,60,40,18,18,24,32,44])
    finance = spec('Financiamiento', question_headers, [],
                   'Selección múltiple: personas que seleccionan / personas con alguna selección en la pregunta. La suma puede superar 100%. Sin respuesta válida incluye vacíos/saltos. Instrumento común 2026, no códigos 2027.', [12,30,20,60,40,18,18,24,32,44])
    opens = spec('Respuestas abiertas', ['Período','Grupo','Código literal','ID respuesta','Respuesta original','Incluida (fx)'], [],
                 'Texto original o tipo no verificado con alta diversidad; sin categorías inferidas. Filtre Incluida=1; uso interno.', [12,30,20,18,90,18])
    source_specs, question_checks, archived_columns = [], [], []
    for response in model['responses']:
        period, headers = response['period'], response['headers']
        timing = re.compile(r'^(?:G\d+Q\d+Time|groupTime\d+|interviewtime)(?:\.|$)',re.I)
        keep = [i for i,header in enumerate(headers) if not timing.match(header)]
        for i,header in enumerate(headers):
            if i not in keep:
                if question_column(header):
                    raise ValueError('No se permite retirar una pregunta del detalle anual.')
                archived_columns.append([f'P{period}',i+1,header,
                                         'Temporizacion tecnica; conservada intacta en 03_RESPALDOS/P'+str(period)+'/01_ENTRADAS'])
        n = len(keep)
        raw = spec(f'Respuestas P{period}', [headers[i] for i in keep] + ['Incluida (fx)', 'Completa incluida (fx)', 'ID normalizado (preparación)'],
                   [[row[i] for i in keep] + [None, None, identifier(row[response['id_column']])] for row in response['rows']],
                   'Preguntas y respuestas originales; tiempos tecnicos en 03_RESPALDOS, inventariados en auditoria. No editar la fuente.')
        raw['stream_source_literals']=True
        for i in range(len(response['rows'])):
            indices = response['selected'].get(i,[])
            raw['formulas'][f'{col(n)}{i+7}'] = '='+selected_profiles(indices)
            raw['formulas'][f'{col(n+1)}{i+7}'] = '='+selected_profiles(indices,'SI')
        end = max(7, len(response['rows'])+6)
        included = ref(raw['name'], n+1, 7, end)
        response_profiles = [selected_profiles(response['selected'].get(i,[]),'SI')
                             for i in range(len(response['rows']))]

        def weighted_count(predicate):
            weights = Counter(expression for i,expression in enumerate(response_profiles)
                              if expression != '0' and predicate(i,response['rows'][i]))
            return '=SUM('+','.join(f'{expression}*{count}' for expression,count in sorted(weights.items()))+')' if weights else '=0'

        for q in response['columns']:
            if q['aggregate']:
                choices = ['Seleccionada'] if q['multiple'] else q['choices']
                valid = (lambda i,row: q['parent_valid'][i]) if q['multiple'] else (lambda i,row: text(row[q['column']]) in choices)
                denominator = weighted_count(valid)
                # Exact source categories are counted once per profile combination.
                # MAX keeps a response attributed to multiple careers counted once.
                target = finance if 'financi' in q['label'].lower() else questions
                first_row = len(target['rows'])+7
                for choice in choices or ['Sin valores observados']:
                    r = len(target['rows']) + 7
                    target['rows'].append([f'P{period}', q['group'], q['code'], q['label'], choice, None, None, None, None, q['status']])
                    matches = (lambda i,row: text(row[q['column']]).lower() in YES) if q['multiple'] else (lambda i,row: text(row[q['column']]) == choice)
                    target['formulas'][f'F{r}'] = weighted_count(matches) if choices else '=0'
                    target['formulas'][f'G{r}'] = denominator if r==first_row else f'=G{first_row}'
                    target['formulas'][f'H{r}'] = f'=IF(G{r}=0,"No aplica",F{r}/G{r})' if q['percent_verified'] else '="Tipo no verificado"'
                    target['formulas'][f'I{r}'] = f'=SUM({included})-G{r}' if r==first_row else f'=I{first_row}'
                    question_checks.append((period,q['column'],choice,target['name'],r))
            else:
                for i,row in enumerate(response['rows']):
                    if text(row[q['column']]).lower() in EMPTY:
                        continue
                    r = len(opens['rows'])+7
                    opens['rows'].append([f'P{period}',q['group'],q['code'],identifier(row[response['id_column']]),row[q['column']],None])
                    opens['formulas'][f'F{r}'] = '='+ref(raw['name'],n+1,i+7)
        source_specs.append(raw)
    for sheet in [questions,finance]:
        sheet['percent_columns'] = [7]
    dictionary = spec('Diccionario', ['Período','Encuesta','Código','Grupo','Texto','Tipo','Comparabilidad','Columna fuente','Inventario'], model['dictionary'],
                      'Equivalencia P1/P2/P3 autorizada para 2026 y condicionada a códigos/textos verificados. Inventario de referencia P2; no acredita IDs ni filtros técnicos de P1/P3. No extrapolar a 2027.', [12,14,22,30,70,28,44,60,40])
    main = [summary,controls,participation,questions,finance,performance,opens,base,profiles,counts,*source_specs,dictionary]
    performance['decimal_columns'] = [4]
    base['decimal_columns'] = [11,21]
    views, panel_sources = dashboards.panels(model,main,spec,ref,col,criterion)
    main = [*views,*main,*panel_sources]
    presence = []
    for code in sorted({row[2] for row in model['dictionary']}):
        present = {row[0]:row[4] for row in model['dictionary'] if row[2]==code}
        presence.append([code,*['Presente' if f'P{p}' in present else 'Columna no disponible' for p in (1,2,3)],
                         'Textos distintos' if len(set(present.values()))>1 else 'Texto coincidente; no prueba equivalencia de opciones o filtros'])
    audit = [spec('Fuentes',['Período','Encuesta','Rol','Ruta','SHA-256'],model['sources']),
             spec('Columnas en respaldo',['Período','Columna original','Encabezado original','Conservación'],archived_columns,
                  'Solo tiempos tecnicos no usados por indicadores. Ninguna pregunta o respuesta sustantiva se retira.'),
             spec('Incidencias',['Período','Persona','Tipo','Detalle'],model['issues']),
             spec('Comparabilidad',dictionary['headers'],model['dictionary'],dictionary['note']),
             spec('Presencia de códigos',['Código literal','P1','P2','P3','Resultado'],presence,
                  'Se compara disponibilidad de columnas, no se infiere aplicación ni equivalencia de preguntas.'),
             spec('Cobertura',['Dimensión','Registros sin clasificación','Tratamiento'],[
                 ['Modalidad',sum(r[3]=='No determinada' for r in records),'Regla institucional V=Presencial; otras jornadas por catálogo. No usar modalidad escolar. Discrepancias en Enriquecimiento anual.'],
                 ['Sector',sum(r[2]=='Sin clasificar' for r in records),'Conservar dentro del universo y exponer en filtros'],
                 ['Promedio AC',sum(r[11] is None for r in records),'No convertir vacíos a cero ni incluirlos en el promedio']]),
             spec('Conciliación',['Concepto','Cantidad'],[[k,v] for k,v in zip(
                 ['Registros de matrícula','Personas únicas','Completas','Incompletas sin completa','Sin respuesta'],reference_counts(model))],
                 'Control independiente sin filtros. Los datos del libro principal cambian con la selección; este control representa el corte completo.')]
    audit.extend([
        spec('Enriquecimiento anual', ['Período', 'CODCLI', 'CODCARPR', 'Sector previo', 'Sector anual', 'Modalidad',
             'Estado código SIES', 'Candidatos SIES', 'Regla modalidad', 'Promedio previo heredado', 'Promedio notas 2026',
             'Notas válidas', 'Períodos académicos con notas', 'Estado rendimiento',
             'Modalidad según catálogo', 'Regla del catálogo', 'Discrepancia institucional con catálogo'], model.get('enrichment_audit', []),
             'Detalle privado por matrícula. Conserva el valor previo; enriquecimiento exclusivo de la capa anual. Sin cambios a fuentes, elegibilidad o cierre FCU.'),
        spec('Calidad notas', ['Período ingreso', 'Control de fuente académica', 'Filas'], model.get('grade_quality', []),
             'Control sobre cada fuente completa: si P1/P2/P3 comparten archivo, no sumar sus controles. Duplicados exactos se colapsan; intentos contradictorios se excluyen.')])
    cases = filter_tests(model)
    for case in cases:
        selection = dict(case['selection'])
        case['expected_cells'] = (expected_question_cells(model,question_checks,selection)
                                  + expected_dimension_cells(model,participation,performance,selection)
                                  + dashboards.expected_cells(model,selection))
    return {'main':main, 'audit':audit, 'year':model['year'], 'policy':POLICY,
            'expected':reference_counts(model), 'expected_cells':expected_question_cells(model,question_checks)
                + expected_dimension_cells(model,participation,performance) + dashboards.expected_cells(model),
            'filter_tests':cases}


def expected_dimension_cells(model, participation, performance, selection=None):
    selection = selection or {}
    selected = [r for r in model['records'] if all(r[d+1] in values for d,values in selection.items())]
    checks = []
    for i, row in enumerate(participation['rows'], 7):
        dimension, value = row[:2]
        segment = [r for r in selected if dimension == 'Anual'
                   or (dimension == 'Cruce' and r[1:4] == row[7:10])
                   or (dimension == 'Sector período' and r[1:3] == json.loads(value))
                   or (dimension == 'Sector carrera' and [r[2],r[5]] == json.loads(value))
                   or (dimension in DIMENSIONS and r[DIMENSIONS.index(dimension)+1] == value)]
        counts = reference_counts({'records':segment})[1:]
        checks.extend(('Participación',f'{col(j+2)}{i}',v) for j,v in enumerate(counts))
        checks.append(('Participación',f'G{i}',counts[1]/counts[0] if counts[0] else 'No aplica'))
    for i, row in enumerate(performance['rows'], 7):
        values = [r[11] for r in selected if (r[1],r[2],r[12]) == tuple(row[:3]) and r[11] is not None]
        checks.extend([('Rendimiento',f'D{i}',len(values)),
                       ('Rendimiento',f'E{i}',sum(values)/len(values) if values else 'Sin notas disponibles al corte')])
    return checks


def expected_question_cells(model, checks, selection=None):
    selection = selection or {}
    selected = {i for i,row in enumerate(model['records']) if row[9]=='SI'
                and all(row[int(d)+1] in values for d,values in selection.items())}
    calculated = {}
    for response in model['responses']:
        complete_rows = [r for r,indices in response['selected'].items() if selected.intersection(indices)]
        for question in response['columns']:
            if not question['aggregate']:
                continue
            values = [text(response['rows'][r][question['column']]) for r in complete_rows]
            if question['multiple']:
                count = sum(v.lower() in YES for v in values)
                valid = sum(question['parent_valid'][r] for r in complete_rows)
                calculated[(response['period'],question['column'])] = {'Seleccionada':count},valid,len(values)-valid,True
            else:
                valid = [value for value in values if value.lower() not in EMPTY]
                calculated[(response['period'],question['column'])] = Counter(valid),len(valid),len(values)-len(valid),question['percent_verified']
    result = []
    for period,column,choice,sheet,row in checks:
        counts,valid,missing,verified = calculated[(period,column)]
        numerator = counts.get(choice,0)
        percent = ('No aplica' if valid==0 else numerator/valid) if verified else 'Tipo no verificado'
        result.extend([(sheet,f'F{row}',numerator),(sheet,f'G{row}',valid),(sheet,f'H{row}',percent),(sheet,f'I{row}',missing)])
    return result


def filter_tests(model):
    tests = []
    for mask in range(1,8):
        selected = {0:{f'P{p}' for p in (1,2,3) if mask & (1 << (p-1))}}
        tests.append({'selection': [[i, sorted(v)] for i,v in selected.items()],
                      'expected': reference_counts(model,selected)})
    for d in range(1,6):
        for value in sorted({r[d+1] for r in model['records']}):
            selected = {d:{value},0:{'P1','P3'}}
            tests.append({'selection': [[i,sorted(v)] for i,v in selected.items()], 'expected': reference_counts(model,selected)})
    tests.append({'selection': [[0,[]]], 'expected':[0,0,0,0,0]})
    return tests


def renderer_runtime():
    runtime = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node = Path(os.environ.get('CARACTERIZACION_ANUAL_NODE', str(runtime/('node.exe' if os.name=='nt' else 'bin/node'))))
    modules = Path(os.environ.get('CARACTERIZACION_ANUAL_NODE_MODULES', str(runtime/'node_modules')))
    if not node.is_file() or not (modules/'@oai/artifact-tool').exists():
        raise ValueError('Falta el generador Excel. Configurar CARACTERIZACION_ANUAL_NODE y CARACTERIZACION_ANUAL_NODE_MODULES antes de ejecutar el anual.')
    return node, modules


def render_payload(payload_path, output_dir, verify_filters=True):
    payload=json.loads(Path(payload_path).read_text())
    validate_source_dependencies(payload)
    node, modules = renderer_runtime()
    builder = ROOT/'scripts/exportar_analisis_anual.mjs'
    with tempfile.TemporaryDirectory(prefix='anual_excel_') as temp:
        temp = Path(temp)
        if os.name == 'nt':
            subprocess.run(['cmd','/c','mklink','/J',str(temp/'node_modules'),str(modules)],check=True,capture_output=True)
        else:
            (temp/'node_modules').symlink_to(modules,target_is_directory=True)
        shutil.copy2(builder,temp/builder.name)
        shutil.copy2(ROOT/'scripts/renderizar_paneles_anual.mjs',temp/'renderizar_paneles_anual.mjs')
        # Export needs additional serialization memory beyond the calculated workbook.
        command = [str(node),'--max-old-space-size=6144','--expose-gc',str(temp/builder.name),str(payload_path),str(output_dir)]
        if verify_filters:
            command.append('--verify-filters')
        audit = Path(output_dir)/'04_AUDITORIA'
        audit.mkdir(parents=True,exist_ok=True)
        with (audit/'generacion_excel.log').open('w',encoding='utf-8') as log:
            out_of_memory=False
            with subprocess.Popen(command,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,
                                  text=True,encoding='utf-8',errors='replace') as process:
                for line in process.stdout:
                    out_of_memory=out_of_memory or 'heap out of memory' in line.lower()
                    print(line,end='',flush=True)
                    log.write(line)
                    log.flush()
                status = process.wait()
            if status:
                if out_of_memory:
                    raise ValueError(f'FALLO_MEMORIA_EXCEL: no se completo ni publico la entrega. P1/P2/P3 se conservan; no repetirlos. Diagnostico: {audit/"generacion_excel.log"}')
                raise subprocess.CalledProcessError(status,command)
    workbook=Path(output_dir)/'02_SALIDAS'/f'Caracterizacion_anual_{payload["year"]}.xlsx'
    add_dashboard_navigation(workbook)
    complete_source_sheets(workbook,payload)
    print('[EXCEL] Respuestas originales incorporadas por filas al mismo libro; pendiente validacion final.',flush=True)


def add_dashboard_navigation(path):
    """Native links only: preserve calculated cells/styles/charts exported by Artifact Tool."""
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel = 'http://schemas.openxmlformats.org/package/2006/relationships'
    rns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    ET.register_namespace('', ns)
    ET.register_namespace('r', rns)
    with zipfile.ZipFile(path) as source:
        relationships={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(source.read('xl/_rels/workbook.xml.rels')).findall(f'{{{rel}}}Relationship')}
        sheets=ET.fromstring(source.read('xl/workbook.xml')).find(f'{{{ns}}}sheets')
        edits={}
        for sheet in sheets:
            if sheet.attrib['name'] not in dashboards.NAMES:
                continue
            target=relationships[sheet.attrib[f'{{{rns}}}id']]
            target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
            root=ET.fromstring(source.read(target))
            hyperlinks=ET.Element(f'{{{ns}}}hyperlinks')
            for address,name in [('K1','Filtros')]+[(f'{col(i*2)}4',name) for i,name in enumerate(dashboards.NAMES)]:
                ET.SubElement(hyperlinks,f'{{{ns}}}hyperlink',{'ref':address,'location':f"'{name}'!A1"})
            # ECMA-376: hyperlinks precede print settings and drawings, after validation.
            following={'printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks',
                       'customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing',
                       'legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst'}
            insert=next((i for i,child in enumerate(root) if child.tag.split('}')[-1] in following),len(root))
            root.insert(insert,hyperlinks)
            edits[target]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        temporary=path.with_suffix('.navigation.tmp')
        with zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as output:
            for info in source.infolist():
                output.writestr(info,edits.get(info.filename,source.read(info.filename)))
    temporary.replace(path)


def validate_saved_exports(payload, delivery):
    """Read exported OOXML back; engine checks alone do not certify saved caches."""
    checked = 0
    for section, relative in [('main',f'02_SALIDAS/Caracterizacion_anual_{payload["year"]}.xlsx'),
                              ('audit',f'04_AUDITORIA/Auditoria_caracterizacion_anual_{payload["year"]}.xlsx')]:
        book = load_workbook(delivery/relative,read_only=True,data_only=True)
        try:
            if book.sheetnames != [table['name'] for table in payload[section]]:
                raise ValueError('Las hojas exportadas no coinciden con el modelo anual.')
            checks = defaultdict(dict)
            if section == 'main':
                for sheet, cell, expected in payload['expected_cells']:
                    checks[sheet][cell] = expected
                checks['Resumen'].update({f'B{r}':v for r,v in enumerate(payload['expected'],7)})
            for table in payload[section]:
                sheet = book[table['name']]
                if table.get('dashboard'):
                    # Random access on a read-only sheet reparses its XML each
                    # time. Read the panel once, retaining only checked cells.
                    wanted=set(table['cells']) | set(table['formulas'])
                    exported={cell.coordinate:cell for row in sheet.iter_rows(
                        max_row=table['last_row'],max_col=12) for cell in row
                        if getattr(cell,'coordinate',None) in wanted}
                    for address, expected in table['cells'].items():
                        if address not in exported or exported[address].value != expected:
                            raise ValueError(f'Dato de panel diferente: {sheet.title}!{address}')
                    for address in table['formulas']:
                        if address not in exported:
                            raise ValueError(f'Formula de panel ausente: {sheet.title}!{address}')
                        cell = exported[address]
                        if cell.data_type == 'e' or (cell.value is None and '""' not in table['formulas'][address]):
                            raise ValueError(f'Formula de panel invalida: {sheet.title}!{address}')
                        if cell.alignment.horizontal != 'center' or cell.alignment.vertical != 'center':
                            raise ValueError(f'Alineacion de panel incorrecta: {sheet.title}!{address}')
                        if address in checks[sheet.title]:
                            expected=checks[sheet.title][address]
                            same=(abs(cell.value-expected)<1e-10 if isinstance(expected,float) and isinstance(cell.value,(int,float)) else cell.value==expected)
                            if not same:
                                raise ValueError(f'Conciliacion de panel incorrecta: {sheet.title}!{address}')
                        checked += 1
                    continue
                for row_number,values in enumerate(sheet.iter_rows(min_row=6,max_row=len(table['rows'])+6,
                                              max_col=len(table['headers'])),6):
                    for column,cell in enumerate(values):
                        address = f'{col(column)}{row_number}'
                        if cell.value is not None and (cell.alignment.horizontal != 'center'
                                                       or cell.alignment.vertical != 'center'):
                            raise ValueError(f'Alineacion exportada incorrecta: {sheet.title}!{address}')
                        if address in table['formulas']:
                            if cell.data_type == 'e' or cell.value is None:
                                raise ValueError(f'Cache de formula invalido: {sheet.title}!{address}')
                        else:
                            expected = (table['headers'][column] if row_number==6 else
                                        table['rows'][row_number-7][column])
                            if cell.value != expected and not (cell.value is None and expected == ''):
                                raise ValueError(f'Dato exportado diferente: {sheet.title}!{address}')
                        if address in checks[sheet.title]:
                            expected = checks[sheet.title][address]
                            same = (abs(cell.value-expected)<1e-12 if isinstance(expected,float)
                                    and isinstance(cell.value,(int,float)) else cell.value==expected)
                            if not same:
                                raise ValueError(f'Conciliacion exportada incorrecta: {sheet.title}!{address}')
                            checked += 1
        finally:
            book.close()
    # Verify the exported native objects, not only the in-memory preview.
    book=load_workbook(delivery/f'02_SALIDAS/Caracterizacion_anual_{payload["year"]}.xlsx',data_only=True)
    try:
        for table in payload['main']:
            sheet=book[table['name']]
            if not table.get('dashboard'):
                for column in table.get('decimal_columns',[]):
                    if table['rows'] and sheet.cell(7,column+1).number_format != '0.0':
                        raise ValueError(f'Formato decimal perdido: {sheet.title}')
                continue
            if len(sheet._charts) != len(table['charts']):
                raise ValueError(f'Graficos perdidos: {sheet.title}')
            if sheet['K1'].hyperlink is None or sheet['K1'].hyperlink.location != "'Filtros'!A1":
                raise ValueError(f'Navegacion perdida: {sheet.title}')
            for chart,definition in zip(sheet._charts,table['charts']):
                if chart.type != 'bar' or len(chart.series) != len(definition['series']):
                    raise ValueError(f'Grafico exportado incorrecto: {sheet.title}')
                for series,(_,address,_) in zip(chart.series,definition['series']):
                    expected="'"+sheet.title+"'!"+':'.join(re.sub(r'([A-Z]+)(\d+)',r'$\1$\2',part) for part in address.split(':'))
                    if series.val.numRef.f != expected:
                        raise ValueError(f'Serie no vinculada: {sheet.title}')
    finally:
        book.close()
    evidence = {'savedValuesChecked':checked,'sourceCellsPreserved':True,'formulaCachesValid':True,
                'tableAlignment':'horizontal_center_vertical_center',
                'dashboardSheets':dashboards.NAMES,'nativeChartBindingsChecked':True,'navigationChecked':True,
                'nativeExcelTested':False}
    (delivery/'04_AUDITORIA/validacion_archivos_exportados.json').write_text(
        json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return evidence


def consolidate(plan, base, output_base=None):
    output_base = base if output_base is None else output_base
    resources, academic_cache = enrichment.load_resources(), {}
    model = prepare([collect_period(entry, base/f'P{entry["periodo"]}', resources, academic_cache) for entry in plan])
    payload = make_payload(model)
    inventory = []
    for entry in plan:
        directory = base/f'P{entry["periodo"]}'
        for file in sorted(directory.rglob('*.xlsx')):
            if '01_ENTRADAS' in file.parts or file.name.startswith('~$'):
                continue
            digest = sha(file)
            book = load_workbook(file,read_only=True,data_only=True)
            for sheet in book:
                inventory.append([f'P{entry["periodo"]}',str(file.relative_to(base)),sheet.title,sheet.max_row,sheet.max_column,
                                  'Conservada íntegra en 03_RESPALDOS; agregados originales no se suman',digest])
            book.close()
    payload['audit'].append(spec('Hojas anteriores',['Período','Archivo','Hoja','Filas con encabezados','Columnas','Destino','SHA-256'],inventory))
    delivery = output_base/'entrega'
    delivery.mkdir()
    path = output_base/'modelo_anual_privado.json'
    path.write_text(json.dumps(payload,ensure_ascii=False,allow_nan=False)+'\n',encoding='utf-8')
    render_payload(path,delivery)
    validate_saved_exports(payload,delivery)
    for entry in plan:
        shutil.copytree(base/f'P{entry["periodo"]}',delivery/'03_RESPALDOS'/f'P{entry["periodo"]}')
    publish_period_outputs(base, delivery, plan)
    return delivery


def publish_period_outputs(base, delivery, plan):
    """Expose period deliverables without editing the preserved stage backups."""
    for entry in plan:
        period = f'P{entry["periodo"]}'
        source = base/period/'02_SALIDAS'
        if not source.is_dir() or not any(p.is_file() for p in source.rglob('*')):
            raise ValueError(f'{period}: no hay salidas de etapa para publicar.')
        target = delivery/'02_SALIDAS'/period
        shutil.copytree(source, target)
        expected = {str(p.relative_to(source)): sha(p) for p in source.rglob('*') if p.is_file()}
        actual = {str(p.relative_to(target)): sha(p) for p in target.rglob('*') if p.is_file()}
        if expected != actual:
            raise ValueError(f'{period}: las salidas publicadas difieren de la etapa original.')
