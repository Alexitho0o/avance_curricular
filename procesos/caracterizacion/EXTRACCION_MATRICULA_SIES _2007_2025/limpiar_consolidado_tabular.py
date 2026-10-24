#!/usr/bin/env python3
"""Regenera el consolidado como libro tabular, transparente y sin graficos."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
ARCHIVO = BASE_DIR / "Consolidado_Retencion_Trazable_Completo.xlsx"
TMP = BASE_DIR / "Consolidado_Retencion_Trazable_Completo.tmp.xlsx"
COLOR_ENCABEZADO = "1F497D"

RENOMBRES_HOJAS = {
    "Retención IPSS": "Retención IPSS",
    "Matriz_Retencion_Carrera_Coh": "Matriz Retencion Carrera",
    "Matriz_Evolucion_Anual_IES": "Matriz Evolucion Institucional",
    "Matriz_Brecha_Metodologica": "Matriz Brecha Metodologica",
    "Conciliacion_Oficial_162": "Conciliacion Oficial SIES",
    "Ranking_Carreras_2024": "Ranking Carreras 2024",
    "Cuadratura_Anual": "Cuadratura Anual",
    "Explicacion_Brecha_2024": "Explicacion Brecha 2024",
    "Notas_Metodologicas_SIES": "Notas Metodologicas SIES",
}

HEADERS = {
    "mrun": "Identificador Encriptado del Estudiante",
    "cod_inst": "Codigo de Institucion",
    "nomb_inst": "Nombre de Institucion",
    "cat_periodo": "Ano del Proceso de Matricula",
    "anio_cohorte": "Ano de Cohorte",
    "anio_seguimiento": "Ano de Seguimiento",
    "anio_ing_carr_act": "Ano de Ingreso a la Carrera Actual",
    "cod_carrera": "Codigo de Carrera",
    "nomb_carrera": "Nombre de Carrera",
    "codigo_unico": "Codigo Unico de Programa",
    "cod_sede": "Codigo de Sede",
    "sede": "Nombre de Sede",
    "nomb_sede": "Nombre de Sede",
    "nivel_global": "Nivel Global de Formacion",
    "nivel_carrera_1": "Nivel Especifico de Carrera 1",
    "nivel_carrera_2": "Nivel Especifico de Carrera 2",
    "modalidad": "Modalidad de Estudio",
    "jornada": "Jornada Academica",
    "tipo_plan_carr": "Tipo de Plan de Carrera",
    "dur_estudio_carr": "Duracion de Estudio de Carrera",
    "dur_proceso_tit": "Duracion del Proceso de Titulacion",
    "dur_total_carr": "Duracion Total de Carrera",
    "retenido_inst": "Estudiantes Retenidos Institucional",
    "retenido_carrera": "Estudiantes Retenidos en la Carrera",
    "retenido_programa": "Estudiantes Retenidos en el Programa",
    "tasa_retencion_institucional_pct": "Tasa de Retencion Institucional Porcentual",
    "tasa_retencion_carrera_pct": "Tasa de Retencion en la Carrera Porcentual",
    "tasa_retencion_programa_pct": "Tasa de Retencion en el Programa Porcentual",
    "tamano_cohorte": "Tamano de Cohorte",
    "retenidos_institucional": "Estudiantes Retenidos Institucional",
    "retenidos_carrera": "Estudiantes Retenidos en la Carrera",
    "retenidos_programa": "Estudiantes Retenidos en el Programa",
    "tasa_oficial_pct": "Tasa Oficial SIES Porcentual",
    "tasa_bruta_local_pct": "Tasa Bruta Local Porcentual",
    "tasa_recalculada_pct": "Tasa Recalculada Porcentual",
    "estandar_institucional_pct": "Estandar Institucional Porcentual",
    "delta_interno_pct": "Diferencia Interna en Puntos Porcentuales",
}


def regional(value: object, decimal: bool = False) -> str:
    if pd.isna(value):
        return "-"
    if decimal:
        return f"{float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{int(round(float(value))):,}".replace(",", ".")


def nombre_formal(columna: object) -> str:
    clave = str(columna).strip().lower()
    return HEADERS.get(clave, str(columna).strip())


def preparar_tabla(tabla: pd.DataFrame) -> pd.DataFrame:
    tabla = tabla.copy()
    tabla.columns = [str(c).strip() for c in tabla.columns]
    nombres_originales = {str(c).strip().lower(): c for c in tabla.columns}
    for clave, original in nombres_originales.items():
        if clave in {"anio", "anio_cohorte", "anio_seguimiento", "cat_periodo", "anio_ing_carr_act"}:
            tabla[original] = pd.to_numeric(tabla[original], errors="coerce").round().astype("Int64")
        elif any(token in clave for token in ("tasa", "pct", "delta", "diferencia")):
            valores = pd.to_numeric(tabla[original].astype("string").str.replace(".", "", regex=False).str.replace(",", ".", regex=False), errors="coerce")
            tabla[original] = valores.map(lambda value: regional(value, True) if pd.notna(value) else "-")
        elif any(token in clave for token in ("tamano", "retenidos", "cohorte", "dif_")):
            valores = pd.to_numeric(tabla[original], errors="coerce")
            tabla[original] = valores.map(lambda value: regional(value) if pd.notna(value) else "-")
    tabla.columns = [nombre_formal(c) for c in tabla.columns]
    return tabla.replace({pd.NA: None, float("nan"): None})


def estilo_libro(ruta: Path) -> None:
    libro = load_workbook(ruta)
    relleno = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    for hoja in libro.worksheets:
        hoja.sheet_view.showGridLines = True
        hoja.freeze_panes = "A2"
        hoja._charts = []
        for dimension in hoja.column_dimensions.values():
            dimension.hidden = False
        for celda in hoja[1]:
            celda.fill = relleno
            celda.font = Font(color="FFFFFF", bold=True)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for columna in hoja.columns:
            letra = columna[0].column_letter
            ancho = min(max(len(str(c.value or "")) for c in columna) + 3, 55)
            hoja.column_dimensions[letra].width = max(ancho, 10)
    libro.save(ruta)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
    hojas = pd.read_excel(ARCHIVO, sheet_name=None)
    tablas = {}
    for nombre, tabla in hojas.items():
        destino = RENOMBRES_HOJAS.get(nombre, nombre[:31])
        tablas[destino] = preparar_tabla(tabla)
        logging.info("Preparada %s: %s filas x %s columnas", destino, len(tablas[destino]), len(tablas[destino].columns))
    with pd.ExcelWriter(TMP, engine="openpyxl") as writer:
        for nombre, tabla in tablas.items():
            tabla.to_excel(writer, sheet_name=nombre, index=False)
    estilo_libro(TMP)
    TMP.replace(ARCHIVO)
    logging.info("Archivo tabular generado: %s bytes", ARCHIVO.stat().st_size)
    logging.info("Hojas finales: %s", ", ".join(tablas))


if __name__ == "__main__":
    main()