#!/usr/bin/env python3
"""Genera un Excel trazable de retencion con formulas nativas de Excel."""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
SALIDA_DIR = BASE_DIR / "salida"
FUENTE_ORIGINAL = SALIDA_DIR / "reporte_matricula_ip_san_sebastian_ORIGINAL.csv"
ARCHIVO_SALIDA = SALIDA_DIR / "Consolidado_Retencion_Trazable.xlsx"

ENCABEZADOS_RAW = [
    "MRUN",
    "CAT_PERIODO_COHORTE",
    "COD_INST",
    "NOMB_INST",
    "COD_CARRERA",
    "NOMB_CARRERA",
    "CODIGO_UNICO",
    "RETENIDO_INST",
    "RETENIDO_CARRERA",
    "RETENIDO_PROGRAMA",
]
COLOR_ENCABEZADO = "1F497D"


def cargar_base(ruta: Path) -> pd.DataFrame:
    df = pd.read_csv(ruta, dtype="string", encoding="utf-8-sig")
    df.columns = [col.strip().lower() for col in df.columns]
    requeridas = {
        "cat_periodo", "mrun", "anio_ing_carr_act", "cod_inst", "nomb_inst",
        "cod_carrera", "nomb_carrera", "codigo_unico", "tipo_plan_carr",
    }
    faltantes = sorted(requeridas - set(df.columns))
    if faltantes:
        raise ValueError(f"Faltan columnas requeridas: {faltantes}")
    for col in ("cat_periodo", "anio_ing_carr_act"):
        df[col] = pd.to_numeric(df[col], errors="coerce", downcast="integer")
    df["mrun"] = df["mrun"].str.strip()
    return df[df["mrun"].notna() & df["mrun"].ne("")].copy()


def construir_detalle(df: pd.DataFrame) -> pd.DataFrame:
    """Construye una fila por estudiante/programa de cada cohorte regular."""
    filas = []
    anios = sorted(df["cat_periodo"].dropna().unique())
    for anio in anios:
        siguiente = anio + 1
        cohorte = df[
            (df["cat_periodo"] == anio)
            & (df["anio_ing_carr_act"] == anio)
            & df["tipo_plan_carr"].fillna("").str.casefold().eq("plan regular")
        ].copy()
        futuro = df[df["cat_periodo"] == siguiente].copy()
        if cohorte.empty or futuro.empty:
            continue

        futuro_inst = set(futuro["mrun"].dropna())
        futuro_carrera = set(zip(futuro["mrun"], futuro["cod_carrera"]))
        futuro_programa = set(zip(futuro["mrun"], futuro["codigo_unico"]))
        cohorte = cohorte.drop_duplicates(
            subset=["mrun", "cod_inst", "cod_carrera", "codigo_unico"]
        )
        for fila in cohorte.itertuples(index=False):
            filas.append(
                {
                    "MRUN": fila.mrun,
                    "CAT_PERIODO_COHORTE": int(anio),
                    "COD_INST": fila.cod_inst,
                    "NOMB_INST": fila.nomb_inst,
                    "COD_CARRERA": fila.cod_carrera,
                    "NOMB_CARRERA": fila.nomb_carrera,
                    "CODIGO_UNICO": fila.codigo_unico,
                    "RETENIDO_INST": int(fila.mrun in futuro_inst),
                    "RETENIDO_CARRERA": int(
                        (fila.mrun, fila.cod_carrera) in futuro_carrera
                    ),
                    "RETENIDO_PROGRAMA": int(
                        (fila.mrun, fila.codigo_unico) in futuro_programa
                    ),
                }
            )
    return pd.DataFrame(filas, columns=ENCABEZADOS_RAW)


def aplicar_estilo_encabezado(ws, fila: int, ultima_columna: int) -> None:
    relleno = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    for col in range(1, ultima_columna + 1):
        celda = ws.cell(fila, col)
        celda.fill = relleno
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")


def ajustar_ancho(ws) -> None:
    for columna in ws.columns:
        maximo = max(len(str(celda.value or "")) for celda in columna)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = min(
            max(maximo + 3, 10), 55
        )


def escribir_tabla_institucional(ws, fila_inicial: int, anios: list[int]) -> int:
    ws.cell(fila_inicial, 1, "Tabla A: Retencion Institucional")
    ws.cell(fila_inicial, 1).font = Font(bold=True, size=13)
    encabezado = fila_inicial + 1
    columnas = ["Año Cohorte", "Tamaño Cohorte (N)", "Alumnos Retenidos", "Tasa Retencion (%)"]
    for col, valor in enumerate(columnas, 1):
        ws.cell(encabezado, col, valor)
    aplicar_estilo_encabezado(ws, encabezado, len(columnas))
    for offset, anio in enumerate(anios, 1):
        fila = encabezado + offset
        ws.cell(fila, 1, anio)
        ws.cell(fila, 2, f'=COUNTIFS(Detalle_Raw!B:B,A{fila})')
        ws.cell(fila, 3, f'=COUNTIFS(Detalle_Raw!B:B,A{fila},Detalle_Raw!H:H,1)')
        ws.cell(fila, 4, f'=IFERROR(C{fila}/B{fila},0)')
        ws.cell(fila, 2).number_format = ws.cell(fila, 3).number_format = "#,##0"
        ws.cell(fila, 4).number_format = "0.0%"
    return encabezado + len(anios)


def escribir_tabla_apertura(ws, fila_inicial: int, detalle: pd.DataFrame, nivel: str) -> int:
    es_carrera = nivel == "carrera"
    titulo = "Tabla B: Retencion por Carrera" if es_carrera else "Tabla C: Retencion por Programa"
    ws.cell(fila_inicial, 1, titulo)
    ws.cell(fila_inicial, 1).font = Font(bold=True, size=13)
    encabezado = fila_inicial + 1
    columnas = (
        ["Año Cohorte", "Código Carrera", "Nombre Carrera", "Tamaño Cohorte (N)", "Alumnos Retenidos", "Tasa Retencion (%)"]
        if es_carrera
        else ["Año Cohorte", "Código Unico", "Nombre Carrera", "Tamaño Cohorte (N)", "Alumnos Retenidos", "Tasa Retencion (%)"]
    )
    for col, valor in enumerate(columnas, 1):
        ws.cell(encabezado, col, valor)
    aplicar_estilo_encabezado(ws, encabezado, len(columnas))
    agrupadores = ["CAT_PERIODO_COHORTE", "COD_CARRERA", "NOMB_CARRERA"] if es_carrera else ["CAT_PERIODO_COHORTE", "CODIGO_UNICO", "NOMB_CARRERA"]
    grupos = detalle[agrupadores].drop_duplicates().sort_values(agrupadores)
    for offset, (_, grupo) in enumerate(grupos.iterrows(), 1):
        fila = encabezado + offset
        ws.cell(fila, 1, int(grupo.iloc[0]))
        ws.cell(fila, 2, grupo.iloc[1])
        ws.cell(fila, 3, grupo.iloc[2])
        raw_col = "E" if es_carrera else "G"
        flag_col = "I" if es_carrera else "J"
        ws.cell(fila, 4, f'=COUNTIFS(Detalle_Raw!B:B,A{fila},Detalle_Raw!{raw_col}:{raw_col},B{fila})')
        ws.cell(fila, 5, f'=COUNTIFS(Detalle_Raw!B:B,A{fila},Detalle_Raw!{raw_col}:{raw_col},B{fila},Detalle_Raw!{flag_col}:{flag_col},1)')
        ws.cell(fila, 6, f'=IFERROR(E{fila}/D{fila},0)')
        ws.cell(fila, 4).number_format = ws.cell(fila, 5).number_format = "#,##0"
        ws.cell(fila, 6).number_format = "0.0%"
    return encabezado + len(grupos)


def main() -> None:
    detalle = construir_detalle(cargar_base(FUENTE_ORIGINAL))
    wb = Workbook()
    raw = wb.active
    raw.title = "Detalle_Raw"
    raw.append(ENCABEZADOS_RAW)
    for fila in detalle.itertuples(index=False, name=None):
        raw.append(list(fila))
    aplicar_estilo_encabezado(raw, 1, len(ENCABEZADOS_RAW))
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = raw.dimensions

    calculo = wb.create_sheet("Calculo_Retencion")
    calculo.sheet_view.showGridLines = True
    raw.sheet_view.showGridLines = True
    anios = sorted(detalle["CAT_PERIODO_COHORTE"].unique()) if not detalle.empty else []
    fila = escribir_tabla_institucional(calculo, 1, anios)
    fila = escribir_tabla_apertura(calculo, fila + 3, detalle, "carrera")
    escribir_tabla_apertura(calculo, fila + 3, detalle, "programa")
    calculo.freeze_panes = "A3"
    ajustar_ancho(raw)
    ajustar_ancho(calculo)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(ARCHIVO_SALIDA)
    print(f"Detalle Raw: {len(detalle):,} filas")
    print(f"Archivo generado: {ARCHIVO_SALIDA}")


if __name__ == "__main__":
    main()