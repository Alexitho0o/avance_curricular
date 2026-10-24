#!/usr/bin/env python3
"""Audita las fuentes numericas de los graficos del consolidado XLSX."""

from pathlib import Path
from subprocess import run

from openpyxl import load_workbook


ARCHIVO = Path(__file__).resolve().parent / "Consolidado_Retencion_Trazable_Completo.xlsx"


def referencia(ref):
    if ref is None:
        return "NO DEFINIDA"
    return ref.f or "NO DEFINIDA"


def auditar_hoja(libro, nombre):
    hoja = libro[nombre]
    print(f"\n[{nombre}]")
    print("Columnas visibles:")
    for indice in range(1, hoja.max_column + 1):
        celda = hoja.cell(1, indice)
        if not str(celda.value).startswith("_GRAFICO_"):
            print(f"  {celda.coordinate}: {celda.value!r} | tipo={type(celda.value).__name__} | formato={celda.number_format}")
            if str(celda.value).lower() == "anio" or str(celda.value).lower().startswith("anio_"):
                valores_anio = [hoja.cell(fila, indice).value for fila in range(2, hoja.max_row + 1) if hoja.cell(fila, indice).value is not None]
                enteros = all(isinstance(valor, int) and not isinstance(valor, bool) for valor in valores_anio)
                sin_miles = all("." not in str(valor) and "," not in str(valor) for valor in valores_anio)
                print(f"    verificacion_anio: enteros={enteros} | sin_separador_miles={sin_miles} | muestra={valores_anio[:5]}")

    print("Columnas auxiliares:")
    for indice in range(1, hoja.max_column + 1):
        nombre_columna = str(hoja.cell(1, indice).value)
        if nombre_columna.startswith("_GRAFICO_"):
            letra = hoja.cell(1, indice).column_letter
            valores = [hoja.cell(fila, indice).value for fila in range(2, hoja.max_row + 1)]
            numericos = sum(isinstance(valor, (int, float)) and not isinstance(valor, bool) for valor in valores if valor is not None)
            print(f"  {letra}: {nombre_columna} | ocultas={hoja.column_dimensions[letra].hidden} | numericos={numericos}/{sum(valor is not None for valor in valores)}")

    print(f"Graficos encontrados: {len(hoja._charts)}")
    for indice_grafico, grafico in enumerate(hoja._charts, 1):
        print(f"  Grafico {indice_grafico}: {grafico.title.tx.rich.p[0].r[0].t if grafico.title and grafico.title.tx and grafico.title.tx.rich else 'sin titulo'}")
        for indice_serie, serie in enumerate(grafico.series, 1):
            val_ref = referencia(serie.val.numRef if serie.val else None)
            cat_ref = referencia(serie.cat.numRef if serie.cat and serie.cat.numRef else (serie.cat.strRef if serie.cat else None))
            print(f"    Serie {indice_serie}: valores={val_ref} | categorias={cat_ref}")
            if serie.val and serie.val.numRef:
                formula = serie.val.numRef.f
                rango = formula.split("!")[-1].replace("$", "").split(":")
                datos = hoja[rango[0]:rango[-1]]
                tipos = sorted({type(celda.value).__name__ for fila in datos for celda in fila if celda.value is not None})
                print(f"      tipos_datos_fuente={tipos} | solo_numericos={all(tipo in {'int', 'float'} for tipo in tipos)}")


def main():
    print(f"Archivo auditado: {ARCHIVO}")
    integridad = run(["unzip", "-tq", str(ARCHIVO)], capture_output=True, text=True)
    print(f"Integridad XLSX: {'OK' if integridad.returncode == 0 else 'ERROR'}")
    libro = load_workbook(ARCHIVO, data_only=False)
    auditar_hoja(libro, "Conciliacion_Oficial_162")
    auditar_hoja(libro, "Ranking_Carreras_2024")


if __name__ == "__main__":
    main()