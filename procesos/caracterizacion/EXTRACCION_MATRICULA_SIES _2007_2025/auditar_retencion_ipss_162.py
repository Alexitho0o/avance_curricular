#!/usr/bin/env python3
"""Audita totales IPSS 162 y compara retencion por carrera contra el pais."""

from __future__ import annotations

import argparse
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
IPSS_SOURCE = BASE_DIR / "salida" / "reporte_matricula_ip_san_sebastian_ORIGINAL.csv"
OUTPUT = BASE_DIR / "salida" / "Auditoria_Retencion_IPSS_162.xlsx"
INSTITUTION = "162"


def numero(value: float, decimals: int = 2) -> str:
    """Formatea numeros con miles punto y decimales coma."""
    if pd.isna(value):
        return "-"
    if decimals == 0:
        return f"{int(round(value)):,}".replace(",", ".")
    return f"{value:,.{decimals}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def configurar_log() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")


def cargar_matricula(ruta: Path) -> pd.DataFrame:
    logging.info("Cargando archivo: %s", ruta)
    if ruta.suffix.lower() == ".csv":
        df = pd.read_csv(ruta, dtype="string", encoding="utf-8-sig")
    elif ruta.suffix.lower() in {".xlsx", ".xls"}:
        df = pd.read_excel(ruta, dtype="string")
    else:
        raise ValueError(f"Formato no soportado: {ruta.suffix}")
    df.columns = [col.strip().lower() for col in df.columns]
    required = {
        "cat_periodo", "mrun", "anio_ing_carr_act", "cod_inst", "nomb_inst",
        "cod_carrera", "nomb_carrera", "codigo_unico", "tipo_plan_carr",
    }
    missing = sorted(required - set(df.columns))
    if missing:
        raise ValueError(f"Faltan columnas requeridas: {missing}")
    for col in ("cat_periodo", "anio_ing_carr_act"):
        df[col] = pd.to_numeric(df[col], errors="coerce", downcast="integer")
    df["mrun"] = df["mrun"].str.strip()
    return df[df["mrun"].notna() & df["mrun"].ne("")].copy()


def preparar_cohortes(df: pd.DataFrame) -> pd.DataFrame:
    """Devuelve una fila por estudiante/carrera de cada cohorte regular."""
    filas = []
    for year in sorted(df["cat_periodo"].dropna().unique()):
        cohort = df[
            (df["cat_periodo"] == year)
            & (df["anio_ing_carr_act"] == year)
            & df["tipo_plan_carr"].fillna("").str.casefold().eq("plan regular")
        ].drop_duplicates(["mrun", "cod_inst", "cod_carrera", "codigo_unico"])
        future = df[df["cat_periodo"] == year + 1]
        if cohort.empty or future.empty:
            continue
        inst_keys = set(zip(future.mrun, future.cod_inst))
        career_keys = set(zip(future.mrun, future.cod_carrera))
        program_keys = set(zip(future.mrun, future.codigo_unico))
        for row in cohort.itertuples(index=False):
            filas.append(
                {
                    "anio_cohorte": int(year),
                    "mrun": row.mrun,
                    "cod_inst": row.cod_inst,
                    "nomb_inst": row.nomb_inst,
                    "cod_carrera": row.cod_carrera,
                    "nomb_carrera": row.nomb_carrera,
                    "codigo_unico": row.codigo_unico,
                    "retenido_inst": int((row.mrun, row.cod_inst) in inst_keys),
                    "retenido_carrera": int((row.mrun, row.cod_carrera) in career_keys),
                    "retenido_programa": int((row.mrun, row.codigo_unico) in program_keys),
                }
            )
    result = pd.DataFrame(filas)
    logging.info("Cohortes evaluables: %s filas", numero(len(result), 0))
    return result


def resumen_institucional(cohortes: pd.DataFrame) -> pd.DataFrame:
    return (
        cohortes.groupby("anio_cohorte", as_index=False)
        .agg(tamano_cohorte=("mrun", "nunique"), retenidos=("retenido_inst", "sum"))
        .assign(tasa_retencion_pct=lambda x: 100 * x.retenidos / x.tamano_cohorte)
    )


def resumen_carrera(cohortes: pd.DataFrame) -> pd.DataFrame:
    result = (
        cohortes.groupby(["anio_cohorte", "cod_carrera", "nomb_carrera"], as_index=False)
        .agg(tamano_cohorte=("mrun", "nunique"), retenidos=("retenido_carrera", "sum"))
    )
    result["tasa_retencion_pct"] = 100 * result["retenidos"] / result["tamano_cohorte"]
    return result


def auditar_totales(institucional: pd.DataFrame, carrera: pd.DataFrame) -> pd.DataFrame:
    por_carrera = carrera.groupby("anio_cohorte", as_index=False).agg(
        tamano_carrera=("tamano_cohorte", "sum"), retenidos_carrera=("retenidos", "sum")
    )
    auditoria = institucional.merge(por_carrera, on="anio_cohorte", how="outer").fillna(0)
    auditoria["diferencia_tamano"] = auditoria.tamano_cohorte - auditoria.tamano_carrera
    auditoria["diferencia_retenidos"] = auditoria.retenidos - auditoria.retenidos_carrera
    auditoria["discrepancia"] = (
        auditoria.diferencia_tamano.ne(0) | auditoria.diferencia_retenidos.ne(0)
    )
    hallazgos = auditoria[auditoria.discrepancia]
    logging.info("Discrepancias detectadas: %s", numero(len(hallazgos), 0))
    return auditoria


def benchmark_nacional(nacional: pd.DataFrame, ipss: pd.DataFrame) -> pd.DataFrame:
    logging.info("Calculando benchmark nacional excluyendo COD_INST=162")
    nacional = nacional[nacional.cod_inst.astype("string") != INSTITUTION]
    national_career = resumen_carrera(preparar_cohortes(nacional))
    national_career = national_career.groupby("cod_carrera", as_index=False).agg(
        tasa_nacional_pct=("tasa_retencion_pct", "mean")
    )
    ipss_career = resumen_carrera(ipss).groupby(["cod_carrera", "nomb_carrera"], as_index=False).agg(
        tamano_cohorte=("tamano_cohorte", "sum"), retenidos=("retenidos", "sum")
    )
    ipss_career["tasa_ipss_pct"] = 100 * ipss_career.retenidos / ipss_career.tamano_cohorte
    result = ipss_career.merge(national_career, on="cod_carrera", how="left")
    result["Delta_Nacional_pct"] = result.tasa_ipss_pct - result.tasa_nacional_pct
    return result.sort_values("Delta_Nacional_pct", ascending=False)


def exportar(comparativo: pd.DataFrame, discrepancias: pd.DataFrame, ruta: Path) -> None:
    logging.info("Exportando auditoria: %s", ruta)
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        comparativo.to_excel(writer, sheet_name="Comparativo_Carreras", index=False)
        discrepancias.to_excel(writer, sheet_name="Discrepancias_Totales", index=False)
    book = load_workbook(ruta)
    fill = PatternFill("solid", fgColor="1F497D")
    for sheet in book.worksheets:
        sheet.sheet_view.showGridLines = True
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(max(len(str(cell.value or "")) for cell in column) + 3, 55)
            sheet.column_dimensions[letter].width = width
        headers = {str(cell.value).lower(): cell.column for cell in sheet[1]}
        for header, column in headers.items():
            if any(word in header for word in ("tasa", "delta")):
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, column).number_format = "0,00"
            if any(word in header for word in ("tamano", "retenidos", "diferencia")):
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, column).number_format = "#,##0"
    book.save(ruta)
    logging.info("Exportacion exitosa: %s", ruta)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--nacional", type=Path, required=True, help="CSV/XLSX SIES nacional")
    parser.add_argument("--ipss", type=Path, default=IPSS_SOURCE)
    parser.add_argument("--salida", type=Path, default=OUTPUT)
    args = parser.parse_args()
    configurar_log()
    ipss = cargar_matricula(args.ipss)
    ipss = ipss[ipss.cod_inst.astype("string") == INSTITUTION]
    logging.info("Filas IPSS 162: %s", numero(len(ipss), 0))
    national = cargar_matricula(args.nacional)
    ipss_cohorts = preparar_cohortes(ipss)
    institutional = resumen_institucional(ipss_cohorts)
    career = resumen_carrera(ipss_cohorts)
    discrepancies = auditar_totales(institutional, career)
    comparison = benchmark_nacional(national, ipss_cohorts)
    exportar(comparison, discrepancies, args.salida)


if __name__ == "__main__":
    main()