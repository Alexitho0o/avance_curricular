#!/usr/bin/env python3
"""Agrega matrices prefacturadas al consolidado de retencion IPSS 162."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent
ARCHIVO = BASE_DIR / "Consolidado_Retencion_Trazable_Completo.xlsx"
HOJA_RAW = "Retención IPSS"
HOJAS_NUEVAS = [
    "Matriz_Retencion_Carrera_Coh",
    "Matriz_Evolucion_Anual_IES",
    "Matriz_Brecha_Metodologica",
]
HOJAS_OBSOLETAS = ["Matriz_Retencion_Carrera_Cohorte"]


def formato_entero(value: object) -> str:
    if pd.isna(value):
        return "-"
    return f"{int(round(float(value))):,}".replace(",", ".")


def formato_decimal(value: object) -> str:
    if pd.isna(value):
        return "-"
    return f"{float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def numero_regional(value: object) -> float:
    """Convierte numeros exportados con coma decimal a float."""
    if pd.isna(value) or value == "":
        return float("nan")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    return float(str(value).replace(".", "").replace(",", "."))


def cargar_raw() -> pd.DataFrame:
    logging.info("Leyendo hoja granular: %s", HOJA_RAW)
    raw = pd.read_excel(ARCHIVO, sheet_name=HOJA_RAW)
    raw.columns = [str(c).strip().lower() for c in raw.columns]
    requeridas = [
        "anio_cohorte", "anio_seguimiento", "cod_carrera", "nomb_carrera",
        "codigo_unico", "modalidad", "jornada", "retenido_inst",
        "retenido_carrera", "retenido_programa",
    ]
    faltantes = [col for col in requeridas if col not in raw.columns]
    if faltantes:
        raise ValueError(f"Faltan columnas en Retención IPSS: {faltantes}")
    for col in ("anio_cohorte", "anio_seguimiento"):
        raw[col] = pd.to_numeric(raw[col], errors="coerce").astype("Int64")
    for col in ("retenido_inst", "retenido_carrera", "retenido_programa"):
        raw[col] = raw[col].map(numero_regional).fillna(0).astype("int64")
    logging.info("Registros granulares cargados: %s", formato_entero(len(raw)))
    return raw


def matriz_carrera_cohorte(raw: pd.DataFrame) -> pd.DataFrame:
    claves = ["anio_cohorte", "cod_carrera", "nomb_carrera", "codigo_unico", "modalidad", "jornada"]
    matriz = raw.groupby(claves, dropna=False, as_index=False).agg(
        tamano_cohorte=("anio_cohorte", "size"),
        retenidos_institucional=("retenido_inst", "sum"),
        retenidos_carrera=("retenido_carrera", "sum"),
        retenidos_programa=("retenido_programa", "sum"),
    )
    for nivel in ("institucional", "carrera", "programa"):
        matriz[f"tasa_retencion_{nivel}_pct"] = (
            100 * matriz[f"retenidos_{nivel}"] / matriz["tamano_cohorte"].replace(0, pd.NA)
        ).round(2)
    return matriz.sort_values(claves).reset_index(drop=True)


def matriz_evolucion(raw: pd.DataFrame) -> pd.DataFrame:
    matriz = raw.groupby(["anio_cohorte", "anio_seguimiento"], as_index=False).agg(
        tamano_cohorte=("anio_cohorte", "size"),
        retenidos_institucional=("retenido_inst", "sum"),
        retenidos_carrera=("retenido_carrera", "sum"),
        retenidos_programa=("retenido_programa", "sum"),
    )
    for nivel in ("institucional", "carrera", "programa"):
        matriz[f"tasa_retencion_{nivel}_pct"] = (
            100 * matriz[f"retenidos_{nivel}"] / matriz["tamano_cohorte"].replace(0, pd.NA)
        ).round(2)
    oficial = pd.read_excel(ARCHIVO, sheet_name="Conciliacion_Oficial_162")
    oficial.columns = [str(c).strip().lower() for c in oficial.columns]
    oficial["anio"] = pd.to_numeric(oficial["anio"], errors="coerce").astype("Int64")
    oficial["tasa_oficial_sies_pct"] = oficial["tasa_oficial_pct"].map(numero_regional)
    matriz = matriz.merge(
        oficial[["anio", "tasa_oficial_sies_pct"]],
        left_on="anio_cohorte",
        right_on="anio",
        how="left",
    ).drop(columns=["anio"])
    return matriz.sort_values(["anio_cohorte", "anio_seguimiento"]).reset_index(drop=True)


def matriz_brecha() -> pd.DataFrame:
    brecha = pd.read_excel(ARCHIVO, sheet_name="Explicacion_Brecha_2024")
    brecha.columns = [str(c).strip().lower() for c in brecha.columns]
    for col in ("cohorte_n", "retenidos", "tasa_local_pct"):
        if col in brecha:
            brecha[col] = brecha[col].map(numero_regional)
    oficial = pd.read_excel(ARCHIVO, sheet_name="Conciliacion_Oficial_162")
    oficial.columns = [str(c).strip().lower() for c in oficial.columns]
    oficial_2024 = oficial[pd.to_numeric(oficial["anio"], errors="coerce").eq(2024)]
    tasa_oficial = numero_regional(oficial_2024["tasa_oficial_pct"].iloc[0])
    brecha["tasa_oficial_sies_2024_pct"] = tasa_oficial
    brecha["diferencia_vs_oficial_pp"] = brecha["tasa_local_pct"] - tasa_oficial
    return brecha


def formatear_matriz(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    for col in result.columns:
        if not pd.api.types.is_numeric_dtype(result[col]):
            continue
        if col == "anio" or col.startswith("anio_"):
            result[col] = result[col].astype("Int64")
        elif "tasa" in col or "pct" in col or "diferencia" in col:
            result[col] = result[col].map(formato_decimal)
        else:
            result[col] = result[col].map(formato_entero)
    return result


def dar_formato_excel(ruta: Path) -> None:
    libro = load_workbook(ruta)
    relleno = PatternFill("solid", fgColor="1F497D")
    for nombre in HOJAS_NUEVAS:
        hoja = libro[nombre]
        hoja.sheet_view.showGridLines = True
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for celda in hoja[1]:
            celda.fill = relleno
            celda.font = Font(color="FFFFFF", bold=True)
        for columna in hoja.columns:
            letra = columna[0].column_letter
            ancho = min(max(len(str(c.value or "")) for c in columna) + 3, 55)
            hoja.column_dimensions[letra].width = max(ancho, 10)
    libro.save(ruta)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s")
    raw = cargar_raw()
    matrices = {
        HOJAS_NUEVAS[0]: matriz_carrera_cohorte(raw),
        HOJAS_NUEVAS[1]: matriz_evolucion(raw),
        HOJAS_NUEVAS[2]: matriz_brecha(),
    }
    libro = load_workbook(ARCHIVO)
    for nombre in HOJAS_OBSOLETAS:
        if nombre in libro.sheetnames:
            del libro[nombre]
    libro.save(ARCHIVO)
    with pd.ExcelWriter(ARCHIVO, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        for nombre, matriz in matrices.items():
            formatear_matriz(matriz).to_excel(writer, sheet_name=nombre, index=False)
            logging.info("Matriz %s: %s filas x %s columnas", nombre, formato_entero(len(matriz)), formato_entero(len(matriz.columns)))
    dar_formato_excel(ARCHIVO)
    logging.info("Archivo actualizado: %s bytes", formato_entero(ARCHIVO.stat().st_size))


if __name__ == "__main__":
    main()