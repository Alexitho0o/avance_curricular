"""Add navigation from user screenshots to the existing FCU action matrix."""
from pathlib import Path
from collections import Counter
import csv
import hashlib
import json
import shutil
from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'gobernanza/FCU/2027/preparacion'
FILE = 'Matriz_codigos_LimeSurvey_acciones_FCU_2027.xlsx'
PATH = BASE / 'evidencias' / FILE
GROUPS = [
    ('ANTECEDENTES PERSONALES', 'G01Q07 G02Q02 G00Q91 G00Q02 G01Q03 G00Q92 G07Q04 G01Q74 G01Q75 G01Q05 G00Q93 G01Q06 G01Q14 G01Q15 G01Q82 G01Q17 G00Q94 G01Q30 G01Q31 G01Q33 G01Q35 G01Q36'),
    ('ANTECEDENTES FAMILIARES Y SOCIOECONÓMICOS', 'G02Q76 G02Q08 G02Q77 G01Q78 G01Q79 G01Q80 G02Q81 G01Q09 G02Q89 G01Q10 G02Q88 G01Q90 G00Q11 G01Q13 G00Q12 G02Q15 G02Q17 G01Q18 G01Q19 G02Q95 G02Q20 G02Q21 G02Q22 G02Q23 G01Q24'),
    ('TRAYECTORIA FORMATIVA Y LABORAL', 'G03Q28 G03Q30 G03Q31 G03Q32 G03Q33 G03Q44 G01Q45 G03Q46 G03Q47 G03Q48 G03Q50 G03Q51 G03Q54 G03Q99 G03Q100 G03Q55 G03Q56 G03Q57'),
    ('ANTECEDENTES DE LA INSTITUCIÓN Y PROGRAMA', 'G04Q37 G04Q39 G01Q50 G01Q41 G01Q43 G01Q55 G01Q85 G01Q84'),
    ('ESTILO DE VIDA Y SUSTENTABILIDAD', 'G05Q57 G05Q86 G01Q58 G01Q59 G05Q60 G01Q87'),
    ('BIENESTAR PARA EL PROCESO EDUCATIVO', ''),
    ('DISPOSICIÓN HACIA LA EDUCACIÓN SUPERIOR', 'G07Q71 G01Q72 G01Q73 G02Q74 G02Q75 G03Q76'),
]
POSITIONS = {code: (i, name, j, 'Grupo y posición visibles en captura 2026-10-05, encuesta 141234')
             for i, (name, codes) in enumerate(GROUPS, 1)
             for j, code in enumerate(codes.split(), 1)}
for code in ['G06Q66', 'G01Q67', 'G01Q68', 'G06Q91', 'G01Q69']:
    POSITIONS[code] = (6, GROUPS[5][0], None, 'Grupo propuesto por tema; sección no desplegada en capturas. Confirmar ubicación y orden.')

def main():
    old_hash = hashlib.sha256(PATH.read_bytes()).hexdigest()
    source = load_workbook(PATH, data_only=False)
    data = list(source['Matriz'].values)
    extra = ['Orden de grupo', 'Grupo en LimeSurvey', 'Posición de pregunta en grupo', 'Evidencia de ubicación']
    already = list(data[0][:4]) == extra
    headers = list(data[0][4:] if already else data[0])
    original = [list(r[4:] if already else r) for r in data[1:]]
    backup = ROOT / 'backups/FCU_2027/matriz_antes_orden_grupos' / (old_hash + '.xlsx')
    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        shutil.copy2(PATH, backup)
    expanded = []
    for row in original:
        code = row[2]
        loc = POSITIONS.get(code, (None, 'PENDIENTE DE UBICAR', None, 'Sin identificador único de pregunta en la fila.'))
        if code == 'G00Q91–G00Q94':
            loc = (1, GROUPS[0][0], None, 'Cuatro preguntas visibles en posiciones 3, 6, 11 y 17; fila agrupada, sin posición única.')
        expanded.append([*loc, *row])
    expanded.sort(key=lambda row: (row[0] or 99, row[2] or 999))
    assert Counter(tuple(r[4:]) for r in expanded) == Counter(tuple(r) for r in original)
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Matriz'
    sheet.append(extra + headers)
    for row in expanded:
        sheet.append(row)
    sheet.freeze_panes = 'E2'
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in {'A':18,'B':52,'C':28,'D':65,'E':20,'F':24,'G':32,'H':28,'I':60,'J':74,'K':68,'L':42,'M':72,'N':75}.items():
        sheet.column_dimensions[col].width = width
    wb.save(PATH)
    check = load_workbook(PATH, read_only=True)
    assert check['Matriz'].max_row == len(original) + 1
    assert check['Matriz'].max_column == 14
    digest = hashlib.sha256(PATH.read_bytes()).hexdigest()
    for sidecar in [PATH.with_suffix('.xlsx.sha256'), BASE / 'hashes' / (FILE + '.sha256')]:
        sidecar.write_text(f'{digest}  {FILE}\n', encoding='utf-8')
    # Synchronize existing governance metadata, preserving unrelated records.
    for folder in ['json', 'validaciones']:
        for path in (BASE / folder).glob('*.json'):
            text = path.read_text(encoding='utf-8')
            if old_hash not in text and FILE not in text:
                continue
            text = text.replace(old_hash, digest).replace('81 filas de contenido; 10 columnas', '81 filas de contenido; 14 columnas')
            records = json.loads(text)
            if isinstance(records, list):
                for rec in records:
                    if rec.get('archivo') == FILE:
                        rec['tamano_bytes'] = PATH.stat().st_size
                        if 'tamano_mb' in rec:
                            rec['tamano_mb'] = round(PATH.stat().st_size / 1048576, 4)
                        if 'fecha_hash' in rec:
                            rec['fecha_hash'] = '2026-10-05'
            path.write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for folder in ['tablas', 'validaciones']:
        for path in (BASE / folder).glob('*.tsv'):
            text = path.read_text(encoding='utf-8')
            if old_hash not in text and FILE not in text:
                continue
            text = text.replace(old_hash, digest).replace('81 filas de contenido; 10 columnas', '81 filas de contenido; 14 columnas')
            rows = list(csv.DictReader(text.splitlines(), delimiter='\t'))
            fields = list(rows[0])
            for rec in rows:
                if rec.get('archivo') == FILE:
                    rec['tamano_bytes'] = str(PATH.stat().st_size)
                    if 'tamano_mb' in rec:
                        rec['tamano_mb'] = str(round(PATH.stat().st_size / 1048576, 4))
                    if 'fecha_hash' in rec:
                        rec['fecha_hash'] = '2026-10-05'
            with path.open('w', encoding='utf-8', newline='') as stream:
                writer = csv.DictWriter(stream, fields, delimiter='\t', lineterminator='\n')
                writer.writeheader()
                writer.writerows(rows)
    for path in (BASE / 'docs').glob('*.md'):
        text = path.read_text(encoding='utf-8')
        if old_hash in text:
            path.write_text(text.replace(old_hash, digest), encoding='utf-8')
    print(json.dumps({'rows':len(expanded),'columns':14,'sha256':digest,'visible_position_rows':sum(r[2] is not None for r in expanded)}, ensure_ascii=False))

if __name__ == '__main__':
    main()
