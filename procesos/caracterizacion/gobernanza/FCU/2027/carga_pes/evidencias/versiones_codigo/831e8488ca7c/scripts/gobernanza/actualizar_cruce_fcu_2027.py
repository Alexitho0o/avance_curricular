"""Prepare current-code patches without modifying the historical workbook."""
from collections import Counter
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys

from openpyxl import load_workbook
import auditar_grupos_fcu_2027 as audit


BASE = audit.BASE
SOURCE = BASE / 'evidencias/Matriz_FCU_2027_GUIA_TRAZABLE_20261006.xlsx'


def prepare(report_path, destination):
    report = json.loads(Path(report_path).read_text())
    policy = json.loads(audit.REVIEW_POLICY.read_text())
    active, excluded = audit.partition_review_checks(report['checks'] + report.get('excluded_checks', []), policy)
    active = audit.apply_user_confirmations(active, report['questions'])
    report.update(checks=active, pending_checks=[c for c in active if not audit.effective_ok(c)], excluded_checks=excluded)
    report['review_policy'] = {'path': str(audit.REVIEW_POLICY), 'version': policy['version'],
                              'sha256': hashlib.sha256(audit.REVIEW_POLICY.read_bytes()).hexdigest()}
    report['checked_settings_ready'] = not report['pending_checks'] and report['disposition_blocks_ok'] \
        and not any(report.get(k) for k in ['orphan_conditions', 'orphan_relevance', 'duplicate_subquestion_codes'])
    questions = report['questions']
    origins = audit.question_origins()
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    raw = {r[0]: r[2] for r in workbook['RAW FCU 2027'].iter_rows(min_row=2, values_only=True)}
    patches, guide_rows, cross_rows = [], [], []
    def patch(sheet, row, col, value):
        patches.append({'sheet': sheet, 'cell': f'{col}{row}', 'value': value})
    additions = ['Codigo actual 2027', 'Clave actual 2027', 'Origen del elemento',
                 'Estado observado en exportacion', 'Texto actual 2027',
                 'Cambios comunicados por usuario', 'Accion historica 20261006']
    for col, label in zip(['R', 'S', 'T', 'U', 'V', 'W', 'X'], additions):
        patch('Cruces', 1, col, label)
    for col, label in zip(['J', 'K', 'L', 'M'],
                          ['Estado exportacion', 'Codigo historico 2026', 'Clave historica 2026', 'Origen']):
        patch('Guía para editar', 1, col, label)
    patch('Guía para editar', 1, 'B', 'Codigo actual 2027 (vacio = historico)')
    patch('Guía para editar', 1, 'C', 'Clave actual 2027')
    patch('Guía para editar', 1, 'F', 'Texto actual / historico si retirado')
    digest = hashlib.sha256(''.join(s['sha256'] for s in report['sources']).encode()).hexdigest()[:10]
    snapshot = BASE / 'evidencias' / f'encuesta_141234_20261007_{digest}'
    snapshot.mkdir(exist_ok=True)
    for source in report['sources']:
        original = Path(source['path'])
        target = snapshot / original.name
        if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest() != source['sha256']:
            raise ValueError(f'Conflicting snapshot: {target}')
        if not target.exists():
            shutil.copy2(original, target)
        if hashlib.sha256(target.read_bytes()).hexdigest() != source['sha256']:
            raise ValueError('Source changed during snapshot')
        source['snapshot'] = str(target)
    institution_extras = {('G04Q39', s) for s in ['SQ003', 'SQ004', 'SQ005', 'SQ006', 'SQ007', 'SQ008', 'SQ018']}
    institution_extras |= {('G01Q50', 'SQ006'), ('G01Q41', 'other'), ('G01Q41', 'SQ017')}
    institution_extras |= {('G04Q37', s) for s in ['SQ009', 'SQ010', 'SQ011']}
    reconciled = {('G04Q37', 'SQ008'): 'Otra razón',
                  ('G04Q39', 'SQ012'): 'Por las modalidades y/o jornadas que ofrece',
                  ('G04Q39', 'SQ016'): 'Por la posibilidad de continuidad de estudios que ofrece'}
    external_fields = {'P4': 'TIPO_DOC', 'P4.1': 'NRO_DOC', 'P5': 'FECHA_NAC',
                       'P12.1': 'CALLE', 'P12.2': 'NRO_CALLE', 'P12.3': 'DEPTO_BLOCK',
                       'P12.4': 'REGION', 'P12.5': 'COMUNA', 'P13': 'CEL', 'P14': 'CORREO',
                       'P16': 'HOGAR_CANT', 'P28': 'COD_INS', 'P29': 'COD_CAR'}
    for rownum, row in enumerate(workbook['Cruces'].iter_rows(min_row=2, values_only=True), 2):
        row = list(row[:17])
        code, key, text = row[1], row[2] or '', row[3] or ''
        if code not in origins and code != 'SIN CÓDIGO CONFIRMADO':
            raise ValueError(f'Unknown historical code: {code}')
        origin = origins.get(code, 'FCU')
        current = []
        refs = re.findall(r'DIS2027\.\d+', str(row[7] or ''))
        if refs:
            current = [m for m in report['disposition_mapping'] if m['manual_id'] in refs]
        elif code in questions:
            q = questions[code]
            if not key:
                current = [{'question': code, 'subquestion': '', 'text': q['text']}]
            elif key == 'other' and q['other'] == 'Y':
                current = [{'question': code, 'subquestion': 'other', 'text': 'Otro'}]
            else:
                targets = [text, str(row[13] or ''), str(raw.get(row[7], '') or '')]
                if (code, key) in reconciled:
                    targets.append(reconciled[(code, key)])
                candidates = q['subquestions'] + [dict(code=a['code'], text=a['label']) for a in q['answers']]
                for sub in candidates:
                    if any(audit.norm(sub['text']) == audit.norm(t) for t in targets if t):
                        current.append({'question': code, 'subquestion': sub['code'], 'text': sub['text']})
        current = list({(m['question'], m['subquestion']): m for m in current}.values())
        category = row[11]
        # Some old SQ identifiers were reused for unrelated official options.
        historical_extra = (code, key) in institution_extras and not current
        if (code, key) in institution_extras:
            origin = 'INSTITUCIONAL'
        failures = [c for c in report['pending_checks'] if c['code'] == code] if not key else []
        if current:
            if (code, key) in reconciled:
                origin = 'FCU'
            current_code = ', '.join(dict.fromkeys(m['question'] for m in current))
            current_key = ', '.join(m['subquestion'] for m in current if m['subquestion'])
            current_text = '\n'.join(m['text'] for m in current)
            category = 'CONSERVAR FCU' if origin == 'FCU' else 'CONSERVAR INSTITUCIONAL'
            status = 'PENDIENTE EN EXPORTACION' if failures else 'PRESENTE'
            action = 'Conservar el elemento identificado por contenido y codigo actual.'
            final = current_text
            confirmations = [c for c in report['checks'] if c['code'] == code and c.get('confirmed_by_user')] if not key else []
            if confirmations and not failures:
                status = 'CONFIRMADO POR USUARIO'
                action = 'Conservar. Funcionamiento verificado visualmente por el usuario; no repetir solicitud de exportacion.'
            if failures:
                if origin == 'FCU':
                    category = 'MODIFICAR FCU'
                action = 'Confirmar en una nueva exportacion los cambios comunicados por el usuario.'
                settings = []
                for c in failures:
                    expected = c['expected']
                    if c['setting'] == 'Condicion trabajo vigente':
                        expected = '(G02Q95.NAOK == "AO01" or G02Q95.NAOK == "AO02")'
                    settings.append(f"{c['setting']}: {expected}")
                final += '\n' + '\n'.join(settings)
        elif historical_extra:
            current_code = current_key = current_text = ''
            category = 'RETIRADO INSTITUCIONAL'
            status = 'RETIRO INSTITUCIONAL HISTORICO'
            action = 'Registrar el retiro observado. No atribuirlo a SIES ni modificar la opcion actual que reutiliza este codigo.'
            final = f'Alternativa institucional historica: {text}\nAusente. Recuperarla solo mediante decision institucional y con clave nueva disponible.'
        elif code == 'SIN CÓDIGO CONFIRMADO' and row[5] in {'P1', 'P2', 'P3', 'P7'}:
            current_code = current_key = current_text = ''
            category = 'RETIRADO FCU'
            status = 'RETIRADO DEL MANUAL 2027'
            action = 'Referencia historica. No crear una pregunta 2027 a partir de esta fila.'
            final = f'Elemento historico: {text}\nFuera del instrumento y esquema de carga FCU 2027.'
        elif code == 'SIN CÓDIGO CONFIRMADO' and row[5] == 'P12':
            current_code = current_key = current_text = ''
            category = 'CONSERVAR FCU'
            status = 'ENCABEZADO DEL MANUAL'
            action = 'Conservar como referencia documental. No crear una pregunta adicional.'
            final = 'DOMICILIO ACTUAL: encabezado que agrupa los campos de direccion.'
        elif code == 'SIN CÓDIGO CONFIRMADO' and row[5] in external_fields:
            current_code = current_key = current_text = ''
            category = 'CONSERVAR FCU'
            status = 'FUENTE DE CARGA POR VALIDAR'
            action = 'Resolver la fuente del campo en el adaptador PES 2027. No agregar automaticamente preguntas a LimeSurvey.'
            final = f"{external_fields[row[5]]}: documentar fuente, regla de conversion y validacion de datos."
        elif not key and code not in questions and code in origins:
            current_code = current_key = current_text = ''
            category = 'RETIRADO FCU' if origin == 'FCU' else 'RETIRADO INSTITUCIONAL'
            status = 'RETIRADO'
            action = 'Retiro ya observado. No volver a eliminar ni restaurar por esta fila historica.'
            final = 'No aparece en la exportacion actual.'
            if code == 'G01Q68':
                final += ' El contenido se conserva en G01Q67[SQ012].'
            if code == 'G01Q73':
                final += ' Los 15 items se conservan en G07Q71 y G01Q72.'
        elif row[11] == 'ELIMINAR FCU' and not refs:
            current_code = current_key = current_text = ''
            status = 'RETIRO HISTORICO'
            category = 'RETIRADO FCU'
            action = 'No eliminar un codigo actual por coincidir con esta clave historica. El texto antiguo no fue encontrado.'
            final = 'Elemento historico retirado; consultar el cruce por contenido.'
        else:
            current_code = current_key = current_text = ''
            status = 'CRUCE PENDIENTE'
            category = 'CONSERVAR FCU' if origin == 'FCU' else 'CONSERVAR INSTITUCIONAL'
            action = 'No editar usando el codigo historico. Verificar equivalencia de contenido antes de cualquier cambio.'
            final = str(row[13] or text)
        group = questions[code]['group'] if code in questions else row[14]
        if code == 'SIN CÓDIGO CONFIRMADO':
            group = 'REFERENCIA DEL MANUAL / DATOS DE CARGA PES'
        position = int(questions[code]['question_order']) + 1 if code in questions and current else None
        if refs and current:
            group = questions[current[0]['question']]['group']
            position = int(questions[current[0]['question']]['question_order']) + 1
        for col, value in [('L', category), ('M', action), ('N', final), ('O', group), ('P', position),
                           ('R', current_code), ('S', current_key), ('T', origin), ('U', status),
                           ('V', current_text), ('W', 'Cambios finales comunicados; no confirmados por nueva exportacion' if failures else ''),
                           ('X', row[11])]:
            patch('Cruces', rownum, col, value)
        if row[11] == 'CONSERVAR INSTITUCIONAL' and origin == 'FCU':
            patch('Cruces', rownum, 'K', 'Clasificacion historica corregida: contenido FCU, no institucional.')
        if historical_extra:
            patch('Cruces', rownum, 'K', 'Alternativa institucional: no figura en las listas oficiales 2026/2027. No heredar el origen FCU del padre.')
        for col, formula in [('B', f'=IF(Cruces!R{rownum}="","",Cruces!R{rownum})'),
                             ('C', f'=IF(Cruces!S{rownum}="","",Cruces!S{rownum})'),
                             ('F', f'=IF(Cruces!V{rownum}="",Cruces!D{rownum},Cruces!V{rownum})'),
                             ('J', f'=Cruces!U{rownum}'), ('K', f'=Cruces!B{rownum}'),
                             ('L', f'=IF(Cruces!C{rownum}="","",Cruces!C{rownum})'), ('M', f'=Cruces!T{rownum}')]:
            patches.append({'sheet': 'Guía para editar', 'cell': f'{col}{rownum}', 'formula': formula})
        guide_rows.append({'row': rownum, 'status': status, 'category': category})
        cross_rows.append({'old_question': code, 'old_subquestion': key, 'old_text': text,
                           'current': current, 'origin': origin, 'status': status})
    result = {'source_workbook': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'report': report, 'patches': patches, 'crosswalk': cross_rows,
              'status_counts': dict(Counter(r['status'] for r in guide_rows)),
              'guide_row_count': len(guide_rows), 'snapshot_directory': str(snapshot),
              'visibility_normalizer': str(Path(__file__).with_name('normalizar_visibilidad_guia_fcu.py')),
              'raw_sheets_unchanged': ['RAW FCU 2026', 'RAW FCU 2027', 'RAW Institucionales']}
    target = Path(destination)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    policy_digest = report.get('review_policy', {}).get('sha256', '')[:8]
    confirmation_digest = hashlib.sha256(audit.USER_CONFIRMATIONS.read_bytes()).hexdigest()[:8]
    report_target = BASE / 'json' / f'AUDITORIA_ESTRUCTURA_20261007_{digest}_{policy_digest}_{confirmation_digest}.json'
    report_target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
    partial_map = {'date': report['date'], 'survey': report['survey'], 'sources': report['sources'],
                   'manual_sha256': hashlib.sha256(audit.REFERENCE['MANUALS'][2027].read_bytes()).hexdigest(),
                   'status': 'PARCIAL_NO_PRODUCTIVO', 'pes_ready': False,
                   'scope': report['scope'], 'review_policy': report.get('review_policy'),
                   'disposition': report['disposition_mapping'],
                   'disability': report['disability_mapping']}
    (BASE / 'json' / 'MAPEO_PARCIAL_FCU_2027_20261007.json').write_text(
        json.dumps(partial_map, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'statuses': result['status_counts'], 'patches': len(patches), 'snapshot': str(snapshot)}, ensure_ascii=False))


if __name__ == '__main__':
    prepare(sys.argv[1], sys.argv[2])
