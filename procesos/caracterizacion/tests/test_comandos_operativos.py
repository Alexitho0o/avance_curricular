from pathlib import Path
import csv
import hashlib
import io
import os
import tempfile
import unittest
from contextlib import redirect_stdout, redirect_stderr
from unittest.mock import patch
import sys
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'scripts/gobernanza'))
from resolver_fuentes_operativas import resolve_sources, SOURCE_VARS
import ejecutar_analisis_anual as annual


class OperationalCommandsTests(unittest.TestCase):
    SURVEYS_2026 = {1: '923565', 2: '443299', 3: '967488'}

    def setUp(self):
        # Routing tests use a one-question instrument; the production schema is
        # independently exercised by test_enriquecimiento_anual.
        gate = patch.object(annual, 'verify_annual_instrument')
        gate.start()
        self.addCleanup(gate.stop)

    def sources(self, directory, survey, *, shared=False):
        directory.mkdir(parents=True, exist_ok=True)
        names = [f'DATOSDEALUMNOS_{"11941" if shared else survey}.xlsx', f'tokens_{survey}.csv',
                 f'results-survey{survey}.xlsx', f'PROMEDIOSDEALUMNOS_{"7804" if shared else survey}.xlsx']
        for name in names:
            file = directory / name
            if file.suffix == '.xlsx':
                workbook = Workbook()
                if name.startswith('results-survey'):
                    workbook.active.append(['id. ID de respuesta', 'submitdate. Fecha de envio',
                                            'lastpage. Ultima pagina', 'token. Codigo de acceso',
                                            'G01Q01. Pregunta de prueba'])
                    workbook.active.append(['1', '2026-01-01 10:00:00', '1', '100', 'Si'])
                else:
                    workbook.active.append(['DATOS DE PRUEBA'])
                workbook.save(file)
                workbook.close()
            else:
                with file.open('w', newline='') as stream:
                    csv.writer(stream).writerows([['token', 'firstname', 'lastname', 'email'],
                                                ['100', 'A', 'B', 'qa@example.invalid']])
        return dict(zip(SOURCE_VARS, [str((directory / name).resolve()) for name in names]))

    def annual_sources(self, directory):
        expected = {}
        for period, survey in self.SURVEYS_2026.items():
            sources = self.sources(directory, survey, shared=True)
            for role in ('RUTA_INSCRITOS', 'RUTA_RESULTADOS'):
                old = Path(sources[role])
                separator = '' if role == 'RUTA_RESULTADOS' and period == 2 else '_'
                name = f'{old.stem}{separator}Encuesta de Caracterizacio\u0301n 2026-{period} ({survey}){old.suffix}'
                new = old.with_name(name)
                old.rename(new)
                sources[role] = str(new)
            expected[period] = sources
        return expected

    def split_annual_sources(self, directory):
        expected = self.annual_sources(directory)
        surveys = directory / 'caracterizacion 1_2_3'
        surveys.mkdir()
        for sources in expected.values():
            for role in ('RUTA_INSCRITOS', 'RUTA_RESULTADOS'):
                old = Path(sources[role])
                new = surveys / old.name
                old.rename(new)
                sources[role] = str(new.resolve())
        return expected

    def test_financing_2027_selects_only_expected_survey(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            self.sources(directory, '967488')
            correct = self.sources(directory, '141234')
            selected, sources = resolve_sources(directory, expected_survey='141234', environ={})
            self.assertEqual(selected, '141234')
            self.assertEqual(sources['RUTA_INSCRITOS'], correct['RUTA_INSCRITOS'])

    def test_stale_explicit_2026_paths_cannot_run_as_2027(self):
        with tempfile.TemporaryDirectory() as temp:
            sources = self.sources(Path(temp), '967488')
            with self.assertRaisesRegex(ValueError, 'otra encuesta'):
                resolve_sources(temp, expected_survey='141234', environ=sources)

    def test_missing_pair_and_ambiguous_period_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            with self.assertRaises(ValueError):
                resolve_sources(directory, environ={})
            self.sources(directory, '100')
            self.sources(directory, '200')
            with self.assertRaises(ValueError):
                resolve_sources(directory, environ={})

    def test_empty_results_are_blocked_before_pipeline(self):
        with tempfile.TemporaryDirectory() as temp:
            sources = self.sources(Path(temp), '141234')
            Path(sources['RUTA_RESULTADOS']).write_bytes(b'')
            with self.assertRaisesRegex(ValueError, r'vacio \(0 bytes\)'):
                resolve_sources(temp, expected_survey='141234', environ={})

    def test_text_renamed_xlsx_is_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            sources = self.sources(Path(temp), '141234')
            Path(sources['RUTA_RESULTADOS']).write_text('not an Excel file')
            with self.assertRaisesRegex(ValueError, 'Excel invalido'):
                resolve_sources(temp, expected_survey='141234', environ={})

    def test_annual_keeps_each_period_sources_and_ignores_global_overrides(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base)
            env = {'CARACTERIZACION_FUENTES_ANUAL':str(base), **expected[2]}
            plan = annual.build_plan(2026, env)
            self.assertEqual([p['periodo'] for p in plan], [1, 2, 3])
            self.assertEqual([p['survey_id'] for p in plan], list(self.SURVEYS_2026.values()))
            for entry in plan:
                self.assertEqual(entry['sources'], expected[entry['periodo']])
                self.assertEqual(entry['mode'], 'FUENTES_ACTUALES_POR_ENCUESTA')
            self.assertEqual(len({file for entry in plan for file in entry['sources'].values()}), 8)

    def test_annual_default_is_downloads_including_current_p2(self):
        with tempfile.TemporaryDirectory() as temp:
            home = Path(temp)
            expected = self.annual_sources(home / 'Downloads')
            self.sources(home / '.local/fuentes_congeladas/2026/P2', '443299')
            with patch.object(annual.Path, 'home', return_value=home):
                plan = annual.build_plan(2026, {})
            self.assertEqual(plan[1]['sources'], expected[2])
            self.assertTrue(all('Downloads' in file for entry in plan for file in entry['sources'].values()))

    def test_annual_prefers_survey_subdirectory_and_administrative_root(self):
        with tempfile.TemporaryDirectory() as temp:
            home = Path(temp)
            base = home / 'Downloads'
            expected = self.split_annual_sources(base)
            for survey in self.SURVEYS_2026.values():
                self.sources(base, survey, shared=True)
            self.sources(base / 'caracterizacion 1_2_3', '141234')
            with patch.object(annual.Path, 'home', return_value=home):
                plan = annual.build_plan(2026, expected[2])
            self.assertEqual([entry['sources'] for entry in plan], list(expected.values()))
            self.assertEqual(len({file for entry in plan for file in entry['sources'].values()}), 8)

    def test_annual_incomplete_subdirectory_never_falls_back_to_root_surveys(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.split_annual_sources(base)
            self.annual_sources(base)
            Path(expected[2]['RUTA_RESULTADOS']).unlink()
            with self.assertRaisesRegex(ValueError, 'P2:.*caracterizacion 1_2_3:.*RUTA_RESULTADOS'):
                annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_subdirectory_never_supplies_missing_administrative_root(self):
        for role in ('RUTA_MATRICULA', 'RUTA_PROMEDIOS'):
            with self.subTest(role=role), tempfile.TemporaryDirectory() as temp:
                base = Path(temp)
                expected = self.split_annual_sources(base)
                self.sources(base / 'caracterizacion 1_2_3', '141234')
                Path(expected[1][role]).unlink()
                with self.assertRaisesRegex(ValueError, f'falta {role}; debe estar fuera'):
                    annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_explicit_overrides_take_priority_over_survey_subdirectory(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.split_annual_sources(base)
            p2 = self.sources(base / 'explicit P2', '443299')
            p3 = self.sources(base / 'explicit P3', '967488')
            env = {'CARACTERIZACION_FUENTES_ANUAL':str(base),
                   'CARACTERIZACION_INPUT_DIR_P2':str(base / 'explicit P2'),
                   **{f'{var}_P3':file for var, file in p3.items()}}
            Path(expected[3]['RUTA_RESULTADOS']).unlink()
            plan = annual.build_plan(2026, env)
            self.assertEqual([entry['sources'] for entry in plan], [expected[1], p2, p3])

    def test_annual_subdirectory_path_must_be_a_directory(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            self.annual_sources(base)
            (base / 'caracterizacion 1_2_3').write_text('not a directory')
            with self.assertRaisesRegex(ValueError, 'no es una carpeta'):
                annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_period_directory_overrides_support_all_three_periods(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            self.annual_sources(base)
            expected = {p: self.sources(base / f'P{p}', survey)
                        for p, survey in self.SURVEYS_2026.items()}
            env = {'CARACTERIZACION_FUENTES_ANUAL':str(base),
                   **{f'CARACTERIZACION_INPUT_DIR_P{p}':str(base / f'P{p}') for p in expected}}
            plan = annual.build_plan(2026, env)
            self.assertEqual([entry['sources'] for entry in plan], list(expected.values()))

    def test_annual_role_overrides_support_all_three_periods(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            current = self.annual_sources(base)
            override = {p: self.sources(base / f'explicit/P{p}', survey)
                        for p, survey in self.SURVEYS_2026.items()}
            env = {'CARACTERIZACION_FUENTES_ANUAL':str(base),
                   **{f'{var}_P{p}':file for p, sources in override.items() for var, file in sources.items()}}
            plan = annual.build_plan(2026, env)
            self.assertEqual([entry['sources'] for entry in plan], list(override.values()))
            self.assertNotEqual(plan[1]['sources'], current[2])

    def test_annual_rejects_wrong_survey_in_period_overrides(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            self.annual_sources(base)
            wrong = self.sources(base / '2027', '141234')
            for period in self.SURVEYS_2026:
                for role in ('RUTA_INSCRITOS', 'RUTA_RESULTADOS'):
                    with self.subTest(period=period, role=role):
                        env = {'CARACTERIZACION_FUENTES_ANUAL':str(base), f'{role}_P{period}':wrong[role]}
                        with self.assertRaisesRegex(ValueError, f'P{period}:.*otra encuesta'):
                            annual.build_plan(2026, env)

    def test_annual_rejects_pair_from_another_2026_period(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base)
            env = {'CARACTERIZACION_FUENTES_ANUAL':str(base),
                   'RUTA_INSCRITOS_P3':expected[2]['RUTA_INSCRITOS'],
                   'RUTA_RESULTADOS_P3':expected[2]['RUTA_RESULTADOS']}
            with self.assertRaisesRegex(ValueError, 'P3:.*otra encuesta'):
                annual.build_plan(2026, env)

    def test_annual_ignores_foreign_surveys_and_excel_lock_files(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base)
            self.sources(base, '141234', shared=True)
            self.sources(base, '9235650', shared=True)
            (base / '~$DATOSDEALUMNOS_99999.xlsx').write_text('lock')
            for sources in expected.values():
                result = Path(sources['RUTA_RESULTADOS'])
                result.with_name('~$' + result.name).write_text('lock')
            plan = annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})
            self.assertEqual([entry['sources'] for entry in plan], list(expected.values()))

    def test_annual_selects_latest_version_only_within_expected_survey(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            self.annual_sources(base)
            newer = self.sources(base, '443299', shared=True)
            newest_time = max(file.stat().st_mtime_ns for file in base.iterdir()) + 1_000_000_000
            for file in newer.values():
                os.utime(file, ns=(newest_time, newest_time))
            plan = annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})
            self.assertEqual(plan[1]['sources'], newer)

    def test_annual_preflight_blocks_everything_if_one_period_is_missing(self):
        with tempfile.TemporaryDirectory() as temp:
            for period in self.SURVEYS_2026:
                with self.subTest(period=period):
                    base = Path(temp) / f'missing_P{period}'
                    expected = self.annual_sources(base)
                    Path(expected[period]['RUTA_RESULTADOS']).unlink()
                    with self.assertRaisesRegex(ValueError, f'P{period}:.*RUTA_RESULTADOS'):
                        annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_requires_shared_promedios(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base)
            Path(expected[1]['RUTA_PROMEDIOS']).unlink()
            with self.assertRaisesRegex(ValueError, 'RUTA_PROMEDIOS'):
                annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':temp})

    def test_annual_other_year_is_rejected(self):
        with self.assertRaisesRegex(ValueError, '2026 P1/P2/P3'):
            annual.build_plan(2027, {})

    def test_annual_preflight_checks_content_and_accepts_code_only_headers(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base)
            file = expected[2]['RUTA_RESULTADOS']
            workbook = load_workbook(file)
            for cell in workbook.active[1]:
                cell.value = cell.value.split('.')[0]
            workbook.save(file)
            workbook.close()
            plan = annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})
            for entry in plan:
                self.assertEqual(entry['source_validation'], {
                    'tokens_rows':1, 'results_rows':1, 'question_columns':1,
                    'matched_tokens':1, 'coverage_pct':100.0})

    def test_text_only_p2_headers_block_before_starting_p1_without_changing_sources(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.split_annual_sources(base / 'Downloads')
            file = Path(expected[2]['RUTA_RESULTADOS'])
            workbook = load_workbook(file)
            for cell in workbook.active[1]:
                cell.value = cell.value.split('. ', 1)[1]
            workbook.save(file)
            workbook.close()
            before = {path:hashlib.sha256(Path(path).read_bytes()).hexdigest()
                      for sources in expected.values() for path in sources.values()}
            output = io.StringIO()
            with (patch.object(annual, 'ROOT', base),
                  patch.object(annual, 'assert_frozen_reference'),
                  patch.dict(annual.os.environ, {'CARACTERIZACION_FUENTES_ANUAL':str(base / 'Downloads')}, clear=True),
                  patch.object(sys, 'argv', ['ejecutar_analisis_anual.py', '2026']),
                  patch.object(annual.subprocess, 'Popen') as process,
                  patch.object(annual, 'snapshot_sources') as snapshot,
                  redirect_stderr(output)):
                self.assertEqual(annual.main(), 3)
            process.assert_not_called()
            snapshot.assert_not_called()
            self.assertFalse((base / '.local').exists())
            self.assertIn('P2:', output.getvalue())
            self.assertIn('Codigo de pregunta y texto de pregunta', output.getvalue())
            self.assertEqual(before, {path:hashlib.sha256(Path(path).read_bytes()).hexdigest() for path in before})

    def test_annual_rejects_questions_without_codes_even_with_valid_metadata(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            sources = self.annual_sources(base)
            file = sources[2]['RUTA_RESULTADOS']
            workbook = load_workbook(file)
            workbook.active['B1'] = 'submitdate'
            workbook.active['E1'] = 'Pregunta sin codigo'
            workbook.save(file)
            workbook.close()
            with self.assertRaisesRegex(ValueError, 'P2:.*codigos de preguntas'):
                annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_requires_metadata_used_by_pipeline_and_consolidation(self):
        for column, field in [('A1', 'id'), ('B1', 'submitdate'), ('D1', 'token')]:
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                base = Path(temp)
                sources = self.annual_sources(base)
                file = sources[2]['RUTA_RESULTADOS']
                workbook = load_workbook(file)
                workbook.active[column] = 'metadata sin codigo'
                workbook.save(file)
                workbook.close()
                with self.assertRaisesRegex(ValueError, 'P2:'):
                    annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_annual_rejects_pairs_without_real_token_overlap(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            sources = self.annual_sources(base)
            file = sources[3]['RUTA_RESULTADOS']
            workbook = load_workbook(file)
            workbook.active['D2'] = '999'
            workbook.save(file)
            workbook.close()
            with self.assertRaisesRegex(ValueError, 'P3:.*cero coincidencias'):
                annual.build_plan(2026, {'CARACTERIZACION_FUENTES_ANUAL':str(base)})

    def test_only_validate_reports_inputs_without_running_or_creating_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.split_annual_sources(base / 'Downloads')
            output = io.StringIO()
            with (patch.object(annual, 'ROOT', base),
                  patch.object(annual, 'assert_frozen_reference'),
                  patch.dict(annual.os.environ, {'CARACTERIZACION_FUENTES_ANUAL':str(base / 'Downloads')}, clear=True),
                  patch.object(sys, 'argv', ['ejecutar_analisis_anual.py', '2026', '--solo-validar']),
                  patch.object(annual.subprocess, 'Popen') as process,
                  patch.object(annual, 'snapshot_sources') as snapshot,
                  redirect_stdout(output)):
                self.assertEqual(annual.main(), 0)
            process.assert_not_called()
            snapshot.assert_not_called()
            self.assertFalse((base / '.local').exists())
            self.assertIn('PREFLIGHT=OK', output.getvalue())
            self.assertIn('Entradas distintas: 8', output.getvalue())
            for sources in expected.values():
                for file in sources.values():
                    self.assertIn(file, output.getvalue())

    def test_preflight_failure_never_starts_even_p1(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            expected = self.annual_sources(base / 'Downloads')
            Path(expected[3]['RUTA_RESULTADOS']).write_bytes(b'')
            output = io.StringIO()
            with (patch.object(annual, 'ROOT', base),
                  patch.object(annual, 'assert_frozen_reference'),
                  patch.dict(annual.os.environ, {'CARACTERIZACION_FUENTES_ANUAL':str(base / 'Downloads')}, clear=True),
                  patch.object(sys, 'argv', ['ejecutar_analisis_anual.py', '2026']),
                  patch.object(annual.subprocess, 'Popen') as process,
                  patch.object(annual, 'snapshot_sources') as snapshot,
                  redirect_stderr(output)):
                self.assertEqual(annual.main(), 3)
            process.assert_not_called()
            snapshot.assert_not_called()
            self.assertFalse((base / '.local').exists())
            self.assertIn('P3:', output.getvalue())

    def test_full_analysis_clears_early_stop_flags_and_sets_cohort(self):
        entry = {'anio':2026, 'periodo':3, 'sources':dict(zip(SOURCE_VARS,['m','t','r','p']))}
        env = annual.child_environment(entry, {'STOP_AFTER_FINANCIAMIENTO':'1',
            'CARACTERIZACION_STOP_AFTER_FINANCIAMIENTO':'1','COHORTES_ANALISIS':'2027'}, Path('/test/P3'))
        self.assertEqual(env['COHORTES_ANALISIS'], '2026')
        self.assertEqual(env['PERIODO_ANALISIS'], '3')
        self.assertEqual(env['STOP_AFTER_FINANCIAMIENTO'], '0')
        self.assertEqual(env['CARACTERIZACION_STOP_AFTER_FINANCIAMIENTO'], '0')
        self.assertEqual(env['CARACTERIZACION_SKIP_FCU'], '0')
        self.assertEqual(env['SOLO_GENERAR_FCU'], '0')

    def test_snapshot_uses_identical_inputs_and_rejects_changed_source(self):
        with tempfile.TemporaryDirectory() as temp:
            base = Path(temp)
            sources = self.sources(base / 'sources', '100')
            entry = {'periodo':1, 'sources':sources,
                     'source_hashes':{var:hashlib.sha256(Path(file).read_bytes()).hexdigest()
                                      for var, file in sources.items()}}
            run = base / 'run'
            run.mkdir()
            execution = annual.snapshot_sources(entry, run)
            self.assertEqual(set(execution['sources']), set(sources))
            self.assertTrue(all(Path(file).is_file() for file in execution['sources'].values()))
            Path(sources['RUTA_INSCRITOS']).touch()
            entry['source_hashes']['RUTA_INSCRITOS'] = 'bad-hash'
            another = base / 'another'
            another.mkdir()
            with self.assertRaisesRegex(ValueError, 'cambio desde la validacion'):
                annual.snapshot_sources(entry, another)


if __name__ == '__main__':
    unittest.main()
