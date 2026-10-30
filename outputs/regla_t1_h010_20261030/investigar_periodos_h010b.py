"""Investigación agregada de períodos; no homologa ni modifica asignaciones."""
import argparse,collections as C,json,os,sqlite3
from pathlib import Path
from construir import sha,dump,txt,aggregate

def main():
 import openpyxl
 ap=argparse.ArgumentParser();ap.add_argument('--repo',type=Path,required=True);ap.add_argument('--privado',type=Path,required=True);a=ap.parse_args();R=a.repo.resolve();P=a.privado.resolve();O=R/'outputs/regla_t1_h010_20261030';os.umask(0o077)
 sources={}
 def load(p):
  p=Path(p);p=p if p.is_absolute() else R/p;sources[str(p)]=sha(p);return json.loads(p.read_text())
 active=load('data/diccionario_sies_seguimiento/activo.json');desc=load(active['descriptor']);b=Path(desc['base_privada']['ruta']);assert sha(b)==desc['base_privada']['sha256'];sources[str(b)]=sha(b);db=sqlite3.connect(b.as_uri()+'?mode=ro',uri=True);rows=[json.loads(t[0]) for t in db.execute('select payload from observaciones')];db.close()
 d7=load('data/diccionario_sies_seguimiento/v7_20261023_inf35_03_t3/descriptor.json');db=sqlite3.connect(Path(d7['base_privada']['ruta']).as_uri()+'?mode=ro',uri=True);hist={z['observacion_id'] for t in db.execute('select payload from observaciones') if not (z:=json.loads(t[0])).get('codigo_interno_adoptado')};db.close()
 target=[r for r in rows if txt(r['periodo']) in {'4','5'}];historical=[r for r in target if r['observacion_id'] in hist];assert len(historical)==604
 dd=load('data/diccionario_umas/DICCIONARIO_UMAS.json');wanted=C.defaultdict(lambda:C.defaultdict(set));regime=C.defaultdict(set)
 for r in target:
  for f in r['fuentes_promedios']:wanted[Path(f['fuente']).name][f['hoja']].add(int(f['fila']))
 metadata=[]
 for f in dd['fuentes']:
  name=f['archivo']
  if name not in wanted:continue
  p=Path(f['ruta']);assert sha(p)==f['sha256'];sources[str(p)]=f['sha256'];wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
  for sheet,line_ids in wanted[name].items():
   if sheet not in wb.sheetnames:continue
   ws=wb[sheet];it=ws.iter_rows(values_only=True);head=[txt(x) for x in next(it)];ri=head.index('REGIMEN') if 'REGIMEN'in head else None
   metadata.append({'archivo':name,'hoja':sheet,'columnas_temporales':[h for h in head if any(w in h.upper() for w in ['FECHA','INICIO','TERMINO'])],'REGIMEN_FUENTE':'REGIMEN','filas_objetivo':len(line_ids)})
   for n,values in enumerate(it,2):
    if n in line_ids and ri is not None:regime[name,sheet,n].add(txt(values[ri]))
  wb.close()
 per_person=C.defaultdict(set);per_exp=C.defaultdict(set)
 for r in rows:
  per_person[r['persona'],int(r['ano'])].add(txt(r['periodo']));per_exp[r['persona'],txt(r['CODCLI']),int(r['ano'])].add(txt(r['periodo']))
 def item(r):
  regs=set()
  for f in r['fuentes_promedios']:regs.update(regime[Path(f['fuente']).name,f['hoja'],int(f['fila'])])
  return {'anio':int(r['ano']),'periodo':txt(r['periodo']),'carrera':r['carrera'],'plan':r['plan'],'REGIMEN_FUENTE':' | '.join(sorted(regs)) if regs else 'NO_RECUPERADO_E','persona_con_123_mismo_anio':bool(per_person[r['persona'],int(r['ano'])]&{'1','2','3'}),'expediente_con_123_mismo_anio':bool(per_exp[r['persona'],txt(r['CODCLI']),int(r['ano'])]&{'1','2','3'})}
 allitems=[item(r) for r in target];hh=[item(r) for r in historical]
 calendar=next((Path.home()/'Library/CloudStorage').glob('OneDrive-Instituto*/0 BASES DE DATOS/Calendario-Academico-por-Periodo_2026_v3.pdf'));assert sha(calendar)=='2167498aabed4a557a70ac51706488ce99a04dfd84cc855f4a2ced80e6ebdd4f';sources[str(calendar)]=sha(calendar)
 from pypdf import PdfReader
 texts=[p.extract_text() for p in PdfReader(calendar).pages];assert 'Inicio de Clases Primer Trimestre' in texts[1] and 'Inicio de Clases Segundo Trimestre' in texts[4] and 'Inicio de Clases tercer Trimestre' in texts[7]
 docs=['docs/SEGUIMIENTO_RUT_Y_CONVERSION_PERIODOS_20261023.md','docs/ESPECIFICACION_PERIODOS_815_H113_20261027.md','data/diccionario_umas/EQUIVALENCIAS_UMAS.json']
 for p in docs:sources[p]=sha(R/p)
 result={'clase':'B/C; homologación E, decisión D pendiente','universos':{'historico_3049_periodos45':604,'v20_periodos45':len(target)},'historicos_por_anio_periodo_regimen':aggregate(hh,['anio','periodo','REGIMEN_FUENTE']),'historicos_por_carrera_plan':aggregate(hh,['anio','periodo','carrera','plan','REGIMEN_FUENTE']),'v20_por_anio_periodo_regimen':aggregate(allitems,['anio','periodo','REGIMEN_FUENTE']),'coexistencia_historicos':aggregate(hh,['periodo','persona_con_123_mismo_anio','expediente_con_123_mismo_anio']),'personas_historicas_unicas':len({r['persona'] for r in historical}),'personas_historicas_con_123':len({r['persona'] for r in historical if per_person[r['persona'],int(r['ano'])]&{'1','2','3'}}),'esquemas_consultados':metadata,'fechas_periodos45':{'inicio':None,'termino':None,'estado':'E','motivo':'PROMEDIOS no contiene inicio/término. FECHAMATRICULA de DA es evento de matrícula, no calendario del período; no sustituye fechas.','requisito':'Definición UMAS y calendario histórico por año/régimen para PERIODO 4 y 5.'},'calendario_2026_B':[{'periodo_etiqueta':'Primer Trimestre','inicio_clases':'2026-03-23','cierre_desde':'2026-06-17','cierre_hasta':'2026-06-18','paginas':[2,5]},{'periodo_etiqueta':'Segundo Trimestre','inicio_clases':'2026-06-30','cierre_desde':'2026-09-30','cierre_hasta':'2026-10-01','paginas':[5,7]},{'periodo_etiqueta':'Tercer Trimestre','inicio_clases':'2026-10-05','cierre_desde':'2026-12-28','cierre_hasta':'2026-12-31','paginas':[8,9]}],'limite_calendario':'Sólo trimestral 2026; no define los valores nativos 4/5 de otros años/régimen. No acredita verano.','documentos':docs,'propuesta':'Conservar 4/5 nativos. No proponer equivalencia estacional no probada. Si se acredita que son actividad del mismo año, someter a D la herencia del T1 de ese año o primer T1 posterior del mismo expediente/plan, con controles anuales; simular no equivale a aplicar.','efecto_simulado':'EFECTO_HIPOTETICO_PERIODOS45.json; escenario condicionado de mismo año; no homologación','fuentes':[{'ruta':p,'sha256':h} for p,h in sorted(sources.items())]}
 dump(O/'PERIODOS45_H010B.json',result)
 lines=['# Períodos 4/5 · H010b','', '**B/C:** observación de los extractos. **E:** significado y fechas históricas. **D pendiente:** cualquier homologación.','',f"Históricos: 604 contextos; {result['personas_historicas_unicas']} personas distintas, {result['personas_historicas_con_123']} también observadas en 1/2/3 en el mismo año. Son unidades diferentes; no se suman.",'','| Año | Período | REGIMEN_FUENTE | Contextos |','|---|---|---|---|']
 for x in result['historicos_por_anio_periodo_regimen']:lines.append(f"| {x['anio']} | {x['periodo']} | {x['REGIMEN_FUENTE']} | {x['n']} |")
 lines+=['','Carreras y planes desglosados en `PERIODOS45_H010B.json → historicos_por_carrera_plan`; datos observados del campo REGIMEN, no inferidos por nombre.','', '## Fechas y documentación','',result['fechas_periodos45']['motivo'],result['fechas_periodos45']['requisito'],'',f'Calendario institucional B: `{calendar}`; SHA-256 `{sha(calendar)}`. T1: 23-03 a cierre 17/18-06 (pp. 2 y 5); T2: 30-06 a cierre 30-09/01-10 (pp. 5 y 7); T3: 05-10 a cierre 28/31-12 (pp. 8 y 9).',result['limite_calendario'],'','La conversión gobernada de NIV_ACA para régimen trimestral transforma niveles académicos; no define PERIODO 4/5. EQUIVALENCIAS_UMAS exige conservarlos. No se localizó una definición de 4/5 en esos documentos ni en el calendario catalogado.','', '## Propuesta no aplicada','',result['propuesta'],'','Decisión Alexi: aprobar / rechazar / pedir definición UMAS: ____','', 'El escenario de efecto está separado en EFECTO_HIPOTETICO_PERIODOS45.json. Las 604 asignaciones permanecen E.']
 (O/'PERIODOS45_H010B.md').write_text('\n'.join(lines)+'\n');print(json.dumps({'regimenes':aggregate(hh,['REGIMEN_FUENTE']),'personas':result['personas_historicas_unicas'],'con_123':result['personas_historicas_con_123']},ensure_ascii=False))
if __name__=='__main__':main()
