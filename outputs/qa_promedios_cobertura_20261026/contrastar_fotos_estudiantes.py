"""Contrasta grupos discrepantes con fotos de matrícula; sin persistir identificadores."""
import collections,glob,hashlib,json,re
from pathlib import Path
import openpyxl
OUT=Path(__file__).parent; prior=json.loads((OUT/'DIFERENCIAS_ASIGNATURA_AGREGADO.json').read_text()); src={k:Path(v['ruta']) for k,v in prior['fuentes'].items()}
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
def clean(x):
 if x is None:return ''
 if isinstance(x,(int,float)) and int(x)==x:return str(int(x))
 return str(x).strip().upper()
def identity(num,dv):
 a=clean(num).replace('.','').replace(' ','');b=clean(dv)
 if '-' in a:a,b=a.split('-',1)
 return (str(int(a)),b) if a.isdigit() and re.fullmatch('[0-9K]',b) else None
data={}
for label,path in src.items():
 assert sha(path)==prior['fuentes'][label]['sha256'];w=openpyxl.load_workbook(path,read_only=True,data_only=True);it=w['Hoja1'].iter_rows(values_only=True);h={clean(v):i for i,v in enumerate(next(it)) if v is not None};d={}
 for r in it:
  def v(k):return clean(r[h[k]]) if k in h and h[k]<len(r) else ''
  if v('ANO')!='2026' or v('PERIODO') not in ('2','3'):continue
  iden=identity(v('RUT'),v('DIG'))
  if not iden or not v('CODCLI'):continue
  key=(v('PERIODO'),*iden,v('CODCLI'),v('CODCARR'),v('PLAN_DE_ESTUDIO'),v('CODRAMO'),v('CODSECC'));d[key]=(v('NOTA_FINAL'),v('ESTADO_ACADEMICO'),v('ESTADO'),v('DESCRIPCION_ESTADO'))
 data[label]=d;w.close()
a=data['respaldo_ampliado'];b=data['v11_12166'];both=a.keys()&b.keys();aonly=a.keys()-b.keys();bonly=b.keys()-a.keys();groups={k[:7] for k in aonly}&{k[:7] for k in bonly}
section={k for k in aonly if k[:7] in groups and k[0]=='3'};grades={k for k in both if k[0]=='3' and a[k][0]!=b[k][0]};states={k for k in both if a[k][1]!=b[k][1]}
def cases(keys):return {(k[1],k[2],k[3]) for k in keys}
targets={'notas_T3':cases(grades),'seccion_T3':cases(section),'estado_T2':cases(k for k in states if k[0]=='2'),'estado_T3':cases(k for k in states if k[0]=='3')};assert {k:len(v) for k,v in targets.items()}=={'notas_T3':585,'seccion_T3':9,'estado_T2':2,'estado_T3':1}
matches=glob.glob('/Users/alexi/Downloads/cierre/*T_2_13_07_2026.xlsx'); matches3=glob.glob('/Users/alexi/Downloads/cierre/*T3_02_10_2026.xlsx');assert len(matches)==len(matches3)==1
photos=[('T2',Path(matches[0]),'Hoja1'),('T3',Path(matches3[0]),'Base')];results=[]
for period,p,sheet in photos:
 w=openpyxl.load_workbook(p,read_only=True,data_only=True);it=w[sheet].iter_rows(values_only=True);h={clean(x):i for i,x in enumerate(next(it)) if x is not None};assert {'RUT','CODCLI'}.issubset(h);records=collections.defaultdict(list);scanned=0
 for row in it:
  def v(k):return clean(row[h[k]]) if k in h and h[k]<len(row) else ''
  if not v('CODCLI'):continue
  scanned+=1;idn=identity(v('RUT'),'')
  if idn:records[(*idn,v('CODCLI'))].append({'estado':v('ESTADOACADEMICO'),'periodo_matricula':v('PERIODOMATRICULA'),'anio_matricula':v('ANOMATRICULA')})
 groups_out={}
 for label,group in targets.items():
  present=group & records.keys();status=collections.Counter();conflicts=0;matches_current=0
  for key in present:
   rows=records[key];states_here={r['estado'] for r in rows if r['estado']};status['|'.join(sorted(states_here)) if states_here else 'SIN_COLUMNA_O_VALOR']+=1;conflicts+=len(states_here)>1
   matches_current+=any(r['anio_matricula']=='2026' and r['periodo_matricula']==period[-1] for r in rows)
  groups_out[label]={'expedientes_objetivo':len(group),'encontrados_mismo_rut_y_codcli':len(present),'no_localizados_en_hoja':len(group-present),'estados_observados':dict(status),'conflictos_estado':conflicts,'con_matricula_anio_periodo_igual_al_archivo':matches_current}
 results.append({'foto':period,'ruta':str(p),'hoja':sheet,'sha256':sha(p),'filas_con_codcli':scanned,'columnas_estado':'ESTADOACADEMICO' in h,'columnas_seccion':'CODSECC' in h,'columnas_nota':'NOTA_FINAL' in h,'resultados':groups_out});w.close()
grado=collections.Counter((a[k][2],a[k][3],b[k][2],b[k][3]) for k in grades)
out={'clasificacion':'B fotos; C cruces; D fotos de matrícula como referencia para estado cuando lo registran; E ausencias y cobertura','fotos':results,'estados_asignatura_diferencia_nota':[{'estado_respaldo':x,'descripcion_respaldo':y,'estado_v11':z,'descripcion_v11':q,'llaves':n} for (x,y,z,q),n in grado.items()],'limites':'Foto T2 no contiene ESTADOACADEMICO; T3 debe interpretarse según su cobertura y alcance de ingreso. Ausencia en foto no prueba deserción; las fotos no incluyen sección ni nota. El período de actividad académica no prueba corte de matrícula.','bases_denominadores_modificados':False}
(OUT/'CONTRASTE_FOTOS_DISCREPANCIAS_AGREGADO.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');print(json.dumps(out,ensure_ascii=False))
