"""Capa local para v11: último estado observado y discrepancias, sin cambiar v11."""
import collections,datetime,glob,hashlib,json,os,re,sys
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT))
from common.diccionario_sies_seguimiento import DiccionarioSiesSeguimiento,sha256
OUT=Path(__file__).parent;PRIVATE=Path('/Users/alexi/.local/share/ipss/v11_ultimo_estado_20261026');SRC=json.loads((OUT/'DIFERENCIAS_ASIGNATURA_AGREGADO.json').read_text())['fuentes']
def clean(x):
 if x is None:return ''
 if isinstance(x,(int,float)) and int(x)==x:return str(int(x))
 return str(x).strip().upper()
def persona(x,dv):
 s=clean(x).replace('.','').replace(' ','');dv=clean(dv)
 if '-' in s:s,dv=s.split('-',1)
 return str(int(s))+'-'+dv if s.isdigit() and re.fullmatch('[0-9K]',dv) else ''
def ultimo(periods):
 if not periods:return {'ano':None,'periodo':None,'estado':None,'resultado':'SIN_ACTIVIDAD_OBSERVADA'}
 p=max(periods);raw=periods[p];nonempty={x for x in raw if x}
 if len(nonempty)==1 and '' not in raw:return {'ano':p[0],'periodo':p[1],'estado':next(iter(nonempty)),'resultado':'ESTADO_DEL_EXTRACTO'}
 return {'ano':p[0],'periodo':p[1],'estado':None,'resultado':'ESTADO_VACIO_O_CONFLICTIVO'}
def main():
 if PRIVATE.exists():raise SystemExit('No sobrescribir capa ya existente')
 if PRIVATE.resolve().is_relative_to(ROOT.resolve()):raise SystemExit('Datos privados no pueden guardarse en repositorio')
 source_periods=collections.defaultdict(lambda:collections.defaultdict(set));source_history=collections.defaultdict(set);files={}; records={}; invalid=collections.Counter()
 for name in ('v11_12166','respaldo_ampliado'):
  p=Path(SRC[name]['ruta']);assert sha256(p)==SRC[name]['sha256'];w=openpyxl.load_workbook(p,read_only=True,data_only=True);it=w['Hoja1'].iter_rows(values_only=True);h={clean(v):i for i,v in enumerate(next(it)) if v is not None}; rows={}
  for r in it:
   def v(k):return clean(r[h[k]]) if k in h and h[k]<len(r) else ''
   pe=persona(v('RUT'),v('DIG'));cli=v('CODCLI');year=v('ANO');period=v('PERIODO');state=v('ESTADO_ACADEMICO')
   if not pe or not cli or not year.isdigit() or not period.isdigit():invalid[name+'_identidad_o_periodo']+=1;continue
   if name=='v11_12166':source_periods[(pe,cli)][(int(year),int(period))].add(state);source_history[(pe,cli)].add(state)
   if year=='2026' and period in ('2','3'):
    key=(pe,cli,year,period,v('CODCARR'),v('PLAN_DE_ESTUDIO'),v('CODRAMO'),v('CODSECC'));rows[key]=(state,v('NOTA_FINAL'))
  w.close();assert sha256(p)==SRC[name]['sha256'];records[name]=rows;files[name]={'ruta':str(p),'sha256':sha256(p),'hoja':'Hoja1'}
 a=records['respaldo_ampliado'];b=records['v11_12166'];common=a.keys()&b.keys();flags=collections.defaultdict(set)
 for k in common:
  context=k[:4]
  if a[k][0]!=b[k][0]:flags[context].add('ESTADO_EXTRACCIONES_DISTINTO')
  if a[k][1]!=b[k][1]:flags[context].add('NOTA_EXTRACCIONES_DISTINTA')
 ao=a.keys()-b.keys();bo=b.keys()-a.keys();groups={k[:7] for k in ao}&{k[:7] for k in bo}
 for k in ao|bo:
  if k[:7] in groups:flags[k[:4]].add('SECCION_EXTRACCIONES_DISTINTA')
 photos={}; patterns=[('2026','2','/Users/alexi/Downloads/cierre/*T_2_13_07_2026.xlsx','Hoja1'),('2026','3','/Users/alexi/Downloads/cierre/*T3_02_10_2026.xlsx','Base')]
 for year,per,pat,sheet in patterns:
  pths=glob.glob(pat);assert len(pths)==1;p=Path(pths[0]);w=openpyxl.load_workbook(p,read_only=True,data_only=True);it=w[sheet].iter_rows(values_only=True);h={clean(v):i for i,v in enumerate(next(it)) if v is not None};d=collections.defaultdict(set)
  for r in it:
   def v(k):return clean(r[h[k]]) if k in h and h[k]<len(r) else ''
   pe=persona(v('RUT'),'');cli=v('CODCLI')
   if pe and cli:d[(pe,cli)].add(v('ESTADOACADEMICO') if 'ESTADOACADEMICO' in h else '')
  w.close();photos[(year,per)]=d;files['foto_'+year+'_'+per]={'ruta':str(p),'sha256':sha256(p),'hoja':sheet,'tiene_estado':'ESTADOACADEMICO' in h}
 svc=DiccionarioSiesSeguimiento();assert svc.edicion=='v11_20261026_86_y_ambiguedad';dbsha=svc.descriptor['base_privada']['sha256'];out=[];tot=collections.Counter();states=collections.Counter(); observations=collections.Counter()
 for oid,pe,cli,year,per in svc.db.execute('SELECT id,persona,codcli,anio,periodo FROM observaciones ORDER BY id'):
  latest=ultimo(source_periods.get((pe,cli),{}));context=(pe,cli,year,per);photo=photos.get((year,per),{}).get((pe,cli));f=sorted(flags.get(context,set()));
  if (year,per) not in photos:photo_state={'resultado':'FOTO_NO_COMPARADA','estados':[]}
  elif photo is None:photo_state={'resultado':'NO_LOCALIZADO_EN_HOJA','estados':[]}
  elif photo=={''}:photo_state={'resultado':'SIN_COLUMNA_ESTADO','estados':[]}
  else:photo_state={'resultado':'OBSERVADO','estados':sorted(x for x in photo if x)}
  if latest['estado'] and photo_state['resultado']=='OBSERVADO' and latest['estado'] not in photo_state['estados']:f.append('ESTADO_PROMEDIOS_ACTUAL_Y_FOTO_DIFIEREN')
  out.append({'observacion_id':oid,'ultimo_periodo_promedios':latest,'foto_matricula':photo_state,'alertas':sorted(set(f))});tot[latest['resultado']]+=1
  for x in set(f):observations[x]+=1
  if latest['estado']:states[latest['estado']]+=1
 svc.close();assert len(out)==25902
 PRIVATE.mkdir(mode=0o700);fd=os.open(PRIVATE/'CAPA_V11_ULTIMO_ESTADO_Y_ALERTAS.json',os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
 with os.fdopen(fd,'w') as f:json.dump(out,f,ensure_ascii=False)
 count=collections.Counter(ultimo(p)['resultado'] for p in source_periods.values());vary=sum(len({x for x in vals if x})>1 for vals in source_history.values());latest_2026=collections.Counter(ultimo(p)['resultado'] for p in source_periods.values() if ultimo(p)['ano']==2026)
 result={'clasificacion':'D criterio interno; B fuentes; C implementación; E fecha/cobertura y discrepancias','fecha_hora':datetime.datetime.now().astimezone().isoformat(),'edicion_diccionario':'v11_20261026_86_y_ambiguedad','sha256_base_v11':dbsha,'fuentes':files,'conteos_por_observacion_v11':dict(tot),'estados_por_observacion_v11':dict(states),'alertas_por_observacion_no_exclusivas':dict(observations),'total_observaciones':len(out),'expedientes_promedios':len(source_periods),'expedientes_con_mas_de_un_estado_observado':vary,'resultado_ultimo_expediente':dict(count),'ultimo_ano_2026':dict(latest_2026),'filas_no_procesadas':dict(invalid),'capa_privada':{'ruta':str(PRIVATE/'CAPA_V11_ULTIMO_ESTADO_Y_ALERTAS.json'),'sha256':sha256(PRIVATE/'CAPA_V11_ULTIMO_ESTADO_Y_ALERTAS.json')},'limites':'Último año/período observado en PROMEDIOS 12166; estado es del extracto y no se asigna retroactivamente a períodos previos. El corte/frescura y cobertura del archivo no están certificados; cero no equivale vacío. Foto no localizada no implica abandono. No afecta códigos SIES ni denominadores.'}
 (OUT/'RESULTADO_CAPA_V11_ULTIMO_ESTADO.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':main()
