"""Cruza casos discrepantes con reportes históricos; publica solo conteos."""
import collections,hashlib,json,re
from pathlib import Path
import openpyxl
OUT=Path(__file__).parent;prior=json.loads((OUT/'DIFERENCIAS_ASIGNATURA_AGREGADO.json').read_text()); paths={k:Path(v['ruta']) for k,v in prior['fuentes'].items()}
reports=Path('/Users/alexi/Documents/GitHub/CARACTERIZACIÓN/.local/analisis_anual/2026/20261016_002933_892820/entrega/03_RESPALDOS')
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
def norm(x):
 if x is None:return ''
 if isinstance(x,(int,float)) and int(x)==x:return str(int(x))
 return str(x).strip().upper()
data={}
for name,p in paths.items():
 assert sha(p)==prior['fuentes'][name]['sha256'];w=openpyxl.load_workbook(p,read_only=True,data_only=True);it=w['Hoja1'].iter_rows(values_only=True);h={norm(x):i for i,x in enumerate(next(it)) if x is not None}; d={}
 for r in it:
  def v(k):return norm(r[h[k]]) if k in h and h[k]<len(r) else ''
  if v('ANO')!='2026' or v('PERIODO') not in ('2','3'):continue
  k=(v('PERIODO'),v('RUT').replace('.','').lstrip('0') or '0',v('DIG'),v('CODCLI'),v('CODCARR'),v('PLAN_DE_ESTUDIO'),v('CODRAMO'),v('CODSECC'));d[k]=(v('NOTA_FINAL'),v('ESTADO_ACADEMICO'))
 data[name]=d;w.close()
a=data['respaldo_ampliado'];b=data['v11_12166'];both=a.keys()&b.keys();grade={k for k in both if k[0]=='3' and a[k][0]!=b[k][0]};state={k for k in both if a[k][1]!=b[k][1]};aonly=a.keys()-b.keys();bonly=b.keys()-a.keys();groups={k[:7] for k in aonly}&{k[:7] for k in bonly};sections={k for k in aonly if k[:7] in groups};assert len(grade)==653 and len(state)==6 and len(sections)==9
def targets(keys):return {(k[1],k[2],k[3]) for k in keys}
cases={'diferencia_nota':targets(grade),'diferencia_estado':targets(state),'diferencia_seccion':targets(sections)}
def rnorm(v):
 s=norm(v).replace('.','').replace(' ','')
 if '-' not in s:return None
 num,dv=s.split('-',1)
 return (str(int(num)),dv) if num.isdigit() and re.fullmatch('[0-9K]',dv) else None
out={'clasificacion':'B contenidos; C cruce; E fecha efectiva','grupos_objetivo_expedientes':{k:len(v) for k,v in cases.items()},'archivos':[],'limite':'Hojas derivadas de datos académicos, no una bitácora UMAS con fecha/hora del evento; no seleccionan sección ni estado al cierre.'}
for per in ('P2','P3'):
 matches=list((reports/per/'02_SALIDAS').glob('PREGRADO_TRAYECTORIA_CARACTERIZACION*.xlsx'));assert len(matches)==1;p=matches[0];w=openpyxl.load_workbook(p,read_only=True,data_only=True); report={'periodo_proceso':per,'ruta':str(p),'sha256':sha(p),'hojas':{}}
 for name in ('CAMBIOS_ACADEMICOS','TRAYECTORIA_POR_PERIODO'):
  s=w[name];it=s.iter_rows(values_only=True);h={norm(x):i for i,x in enumerate(next(it)) if x is not None};found={k:set() for k in cases};tot=0
  for row in it:
   tot+=1
   def v(k):return row[h[k]] if k in h and h[k]<len(row) else None
   ident=rnorm(v('RUT_NORMALIZADO')) or rnorm(v('RUT'));cli=norm(v('CODCLI'))
   if not ident:continue
   key=(*ident,cli)
   for group,ts in cases.items():
    if key in ts:found[group].add(key)
  report['hojas'][name]={'filas_datos':tot,'expedientes_encontrados_por_grupo':{k:len(v) for k,v in found.items()},'campos_fecha_de_evento':[k for k in h if ('FECHA' in k and 'MATRICULA' not in k and 'RESPUESTA' not in k)]}
 w.close();out['archivos'].append(report)
(OUT/'CONTRASTE_CAMBIOS_TRAYECTORIA_AGREGADO.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');print(json.dumps(out,ensure_ascii=False))
