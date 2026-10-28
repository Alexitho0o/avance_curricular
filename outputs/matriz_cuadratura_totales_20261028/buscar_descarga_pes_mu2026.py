"""Búsqueda dirigida de una descarga PES MU2026: metadatos, sin ejecutar cruces."""
import collections,datetime,glob,hashlib,json,os,re,unicodedata
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
def norm(s):return ''.join(x for x in unicodedata.normalize('NFD',str(s).lower()) if not unicodedata.combining(x))
def sha(p):
    with open(p,'rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def term(s):return bool(re.search(r'descarga|cargado|lo cargado',norm(s)))
catalog_records=[];catalog_scans=[]
def walkmeta(x,source):
    if isinstance(x,dict):
        p=x.get('ruta_consolidada',x.get('ruta'))
        if p and x.get('sha256'):catalog_records.append({'ruta':str(p),'sha256':x['sha256'],'catalogo':source})
        for v in x.values():walkmeta(v,source)
    elif isinstance(x,list):
        for v in x:walkmeta(v,source)
for rel in ('data/catalogo_fuentes/CATALOGO_FUENTES.json','data/catalogo_datos/CATALOGO_DATOS_IPSS.json'):
    walkmeta(json.loads((ROOT/rel).read_text()),rel)
    catalog_scans.append({'ruta':str(ROOT/rel),'sha256':sha(ROOT/rel)})
byhash=collections.defaultdict(list)
for r in catalog_records:byhash[r['sha256']].append(r['ruta'])
strong_catalog=[r for r in catalog_records if '2026' in norm(r['ruta']) and term(Path(r['ruta']).name)]
reference2025=[r for r in catalog_records if '2025' in Path(r['ruta']).name and 'descarga de lo cargado' in norm(Path(r['ruta']).name)]
onedrive=glob.glob('/Users/alexi/Library/CloudStorage/OneDrive-*/Datos DAI*/1.2.1 Datos Institucionales/1.2.1.1 SIES/Enviados/Matrícula Unificada/2026')
if len(onedrive)!=1:raise ValueError('Carpeta de búsqueda OneDrive no única')
scopes=[Path(onedrive[0]),ROOT/'.local_restricted/requerimiento_bettersoft_uplus_2026/gate09_evidencia/matricula_unificada_2026',Path('/Users/alexi/Downloads'),Path('/Users/alexi/Desktop')]
inventory=[];scans=[]
generic={str(Path(onedrive[0])/'MATRÍCULA UNIFICADA 2026.xlsx'),str(Path(onedrive[0])/'SEGUNDA CARGA/ARCHIVOS VARIOS/Matricula Unificada_ ABF.xlsx'),str(Path(onedrive[0])/'PRIMERA CARGA/RESPALDO MATRÍCULA UNIFICADA 2026.xlsx')}
for index,base in enumerate(scopes,2):
    scan={'orden':index,'raiz':str(base),'existe':base.exists(),'archivos':0,'directorios':[],'errores':[],'excluidos_tecnicos':[],'coincidencias_patron_descarga_2026':0}
    for here,dirs,files in os.walk(base,onerror=lambda e:scan['errores'].append(type(e).__name__)):
        scan['directorios'].append(str(Path(here).relative_to(base)))
        for d in list(dirs):
            if d in ('.git','node_modules','__pycache__'):
                dirs.remove(d);scan['excluidos_tecnicos'].append(str(Path(here)/d))
        for name in files:
            scan['archivos']+=1;p=Path(here)/name;n=norm(name);whole=norm(p)
            if ('precarga' in n and '2027' in n) or name.startswith('.'):continue
            strong='2026' in whole and term(name)
            if strong:scan['coincidencias_patron_descarga_2026']+=1
            if index>=4 and not strong:continue
            if p.suffix.lower() not in ('.csv','.xlsx','.xls','.zip','.txt') and not strong:continue
            r={'ruta':str(p),'ambito':str(base),'bytes':p.stat().st_size,'coincide_patron':strong,'clasificacion':'B metadatos; D criterio solicitado; E procedencia no acreditada'}
            if '_orig' in n:r.update(estado='DESCARTADO',motivo='Archivo de envío original; no descarga de lo recibido. Criterio explícito del usuario.')
            elif n.endswith('.csv.txt') or 'problemas' in whole or '/errores/' in whole:r.update(estado='DESCARTADO',motivo='Archivo o carpeta de errores de carga; no nómina SIES recibida.')
            elif any(t in n for t in ('para_subir','pes_ready','consolidada_final','reconstruid','auditoria','reporte_95','kpi_','resumen_ejecutivo','detalle_vigencia','detalle_solo_vigentes')):r.update(estado='DESCARTADO',motivo='Preparación/envío o auditoría interna; no descarga individual PES acreditada.')
            elif str(p) in generic:r.update(estado='E',motivo='Libro de trabajo con hojas internas; no acredita procedencia PES ni fecha de descarga posterior al 11-05-2026.')
            elif strong:r.update(estado='E',motivo='Nombre candidato; requiere evidencia de procedencia PES, fecha y esquema individual.')
            else:r.update(estado='FUERA_DE_CRITERIO',motivo='Nombre o contexto sin evidencia de descarga PES individual posterior al 11-05-2026.')
            if strong or str(p) in generic:
                try:
                    r['sha256']=sha(p);r['coincidencias_hash_en_catalogos']=byhash.get(r['sha256'],[])
                    if p.suffix.lower()=='.xlsx':
                        w=openpyxl.load_workbook(p,read_only=True,data_only=True)
                        r['hojas']=[{'hoja':s.title,'filas_declaradas':s.max_row,'columnas_declaradas':s.max_column} for s in w];w.close()
                        r['filas_y_columnas_nota']='Dimensiones declaradas de hojas; no conteo de registros válidos.'
                except OSError as e:r['lectura_error']=type(e).__name__
            inventory.append(r)
    scans.append(scan)
result={'estado':'NO_LOCALIZADA_DESCARGA_QUE_CUMPLA_CRITERIOS','clase':'B inventario; C búsqueda; D criterios; E procedencia pendiente','fecha':'2026-10-28','contexto':'MU 2026 recibida por SIES; sólo búsqueda; sin AC 2027','criterios':['Descarga desde PES posterior al 11-05-2026','Documento+DV y código único a nivel individual','No envío, error, consolidado o reconstrucción IPSS'],'catalogos':catalog_scans,'catalogo_coincidencias_descarga_2026':strong_catalog,'referencia_descarga_2025':reference2025,'ambitos':scans,'candidatos_y_descartes':inventory,'metodo_mu2026_ejecutado':False,'fuentes_copiadas':False,'catalogos_modificados':False,'limite':'No afirma inexistencia universal. No se buscan precargas AC2027. Carpetas técnicas excluidas se listan; metadatos y nombres no acreditan fecha PES.'}
(OUT/'BUSQUEDA_DESCARGA_PES_MU2026.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
lines=['# Búsqueda de descarga PES MU 2026 · 28-10-2026','','Resultado: no se localizó una descarga que acredite los tres criterios. Método no ejecutado.','', '## Catálogos revisados primero','']
for r in catalog_scans:lines.append(f"- {r['ruta']} · SHA-256 {r['sha256']}")
lines+=['','Coincidencias de nombre descarga/cargado con 2026 en catálogo: '+str(len(strong_catalog))+'. Los hashes de candidatos genéricos se contrastaron con todas las entradas de ambos catálogos.','','## Ámbitos físicos revisados','', '| Raíz | Archivos enumerados | Directorios | Coincidencias de descarga 2026 | Errores |','|---|---:|---:|---:|---:|']
for s in scans:lines.append(f"| {s['raiz']} | {s['archivos']} | {len(s['directorios'])} | {s['coincidencias_patron_descarga_2026']} | {len(s['errores'])} |")
lines+=['','Subcarpetas, exclusiones técnicas y rutas exactas de todos los candidatos: JSON adjunto. No se han leído datos de AC 2027.','','## Candidatos y descartes','', '| Archivo (relativo al ámbito) | Estado | Motivo |','|---|---|---|']
for r in inventory:lines.append(f"| {Path(r['ruta']).relative_to(r['ambito'])} | {r['estado']} | {r['motivo']} |")
lines+=['','## Límite y siguiente evidencia','','Los libros genéricos quedan E: sus hojas internas no acreditan descarga PES ni fecha posterior al 11-05-2026. No se usan sus dimensiones como filas recibidas por SIES. No hay candidato aceptado con fecha de descarga acreditada. Se necesita localizar el archivo descargado o su evidencia PES; no se sustituyó por nuestros envíos. Sin copias nuevas, cambios al catálogo, cruces, commit o push.']
(OUT/'BUSQUEDA_DESCARGA_PES_MU2026.md').write_text('\n'.join(lines)+'\n')
print(json.dumps({'resultado':result['estado'],'ambitos':[{'raiz':s['raiz'],'archivos':s['archivos'],'directorios':len(s['directorios']),'patron':s['coincidencias_patron_descarga_2026'],'errores':len(s['errores'])} for s in scans],'candidatos':len(inventory),'por_estado':dict(collections.Counter(r['estado'] for r in inventory)),'metodo_ejecutado':False},ensure_ascii=False))
