"""Auditoría de dos copias de titulación. Solo lee fuentes; exporta agregados.

No reconstruye pesos desconocidos ni adjudica títulos o homologaciones.
Los identificadores se utilizan exclusivamente en memoria.
"""
import collections
import hashlib
import json
import re
import subprocess
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
def sha(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda: f.read(1048576), b''):
            h.update(b)
    return h.hexdigest()
def txt(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v or '').strip().upper()
def body(v):
    return txt(v).replace('.', '').split('-')[0]
def rows(ws):
    it = ws.iter_rows(values_only=True)
    h = [str(x or '') for x in next(it)]
    return [(i, dict(zip(h, r))) for i, r in enumerate(it, 2)]

private = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
cases = json.loads(private.read_text())
wanted = {x['RUT'].split('-')[0] for x in cases}
assert len(wanted) == 4
prior_path = ROOT / 'outputs/conciliacion_22_ramos_titulados_20261026/EVIDENCIA_AGREGADA.json'
source = json.loads(prior_path.read_text())['fuentes']['respaldo_xlsx']
primary = Path(source['ruta'])
teams_dir = Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Archivos de chat de Microsoft Teams')
candidates = list(teams_dir.glob('1668382*.xlsx'))
assert len(candidates) == 1
teams = candidates[0]
assert sha(primary) == source['sha256']
sources = []
for p in [primary, teams]:
    digest = sha(p)
    w = openpyxl.load_workbook(p, read_only=True, data_only=False)
    inventory = []
    for s in w:
        count = collections.Counter()
        examples = {}
        for row in s.iter_rows():
            for c in row:
                if c.data_type == 'f':
                    count['formulas'] += 1
                    count['referencia_rota'] += '#REF!' in c.value
                    examples.setdefault(c.column_letter, {'celda': c.coordinate, 'formula': c.value})
        inventory.append({'hoja': s.title, 'filas': s.max_row, 'columnas': s.max_column,
                          'conteos': dict(count), 'formulas_ejemplo_por_columna': examples})
    da = dict(rows(w['Datos_alumnos']))
    homolog = dict(rows(w['CUADRO HOMOLOGACIÓN']))
    trace = [(i, r) for i, r in rows(w['Res_Tit_Unificada']) if body(r['NUM_DOCUMENTO']) in wanted]
    assert len(trace) == 4
    counts = collections.Counter()
    for _, r in trace:
        counts['trazas'] += 1
        counts['puntaje_2130'] += r['TRACE_SCORE_CANDIDATO'] == 2130
        pointer = int(r['TRACE_ROWINDEX_DATOS_ALUMNOS'])
        # Probar offsets, sin asumir base de índice del generador desconocido.
        for offset in [-1, 0, 1, 2]:
            d = da.get(pointer + offset, {})
            if txt(d.get('CODCLI')) == txt(r['TRACE_CODCLI_DATOS_ALUMNOS']) and body(d.get('RUT')) == body(r['NUM_DOCUMENTO']):
                counts[f'puntero_da_concordante_offset_{offset}'] += 1
        counts['codcli_y_documento_localizados_en_alguna_fila_da'] += any(
            txt(d.get('CODCLI')) == txt(r['TRACE_CODCLI_DATOS_ALUMNOS'])
            and body(d.get('RUT')) == body(r['NUM_DOCUMENTO']) for d in da.values())
        for offset in [-1, 0, 1, 2]:
            h = homolog.get(int(r['TRACE_HOMO_OBT_ROW']) + offset, {})
            if (txt(h.get('CODIGO_SIES')) == txt(r['TRACE_HOMO_OBT_CODIGO'])
                and txt(h.get('CODCARPR')) == txt(r['TRACE_CODCARPR_DATOS_ALUMNOS'])
                and txt(h.get('JORNADA_DA')) == txt(r['TRACE_JORNADA_DATOS_ALUMNOS'])):
                counts[f'puntero_homologacion_concordante_offset_{offset}'] += 1
    # Las dos columnas auxiliares consultan un solo documento, no CODCLI ni plan.
    lookup_checks = collections.Counter()
    for _, r in rows(w['Matricula Unificada 2026 subir']):
        if body(r['NUM_DOCUMENTO']) not in wanted:
            continue
        lookup_checks['filas_objetivo'] += 1
        exact = [(i, d) for i, d in da.items() if 2 <= i <= 9677 and d.get('RUT') == r['NUM_DOCUMENTO']]
        lookup_checks['con_coincidencia_valor_y_tipo_comparable'] += bool(exact)
        lookup_checks['con_varios_codcli_en_rango'] += len({txt(d['CODCLI']) for i, d in exact}) > 1
        target = [t for i, t in trace if body(t['NUM_DOCUMENTO']) == body(r['NUM_DOCUMENTO'])]
        if exact and len(target) == 1:
            lookup_checks['primera_coincidencia_codcli_igual_traza'] += txt(exact[0][1]['CODCLI']) == txt(target[0]['TRACE_CODCLI_DATOS_ALUMNOS'])
    header_sets = {name: [str(c.value or '') for c in next(w[name].iter_rows())]
                   for name in ['Datos_alumnos', 'Cruce_RUT_Datos_Alumnos', 'CUADRO HOMOLOGACIÓN']}
    w.close()
    assert sha(p) == digest
    sources.append({'ruta': str(p), 'sha256': digest, 'clasificacion': 'B/C',
                    'hojas': inventory, 'validacion_traza': dict(counts),
                    'inspeccion_busqueda_auxiliar_no_recalculo_excel': dict(lookup_checks),
                    'cabeceras_auxiliares': header_sets, 'hash_conservado': True})

worktree = Path('/Users/alexi/Documents/GitHub/_worktrees/avance_curricular_personal_academico_siipa_v4')
command = ['rg', '-l', '--hidden']
for g in ['*.py', '*.sql', '*.mjs', '*.js', '*.ipynb', '!**/.git/**', '!**/node_modules/**', '!**/.venv/**', '!**/venv/**']:
    command += ['-g', g]
command += ['TRACE_SCORE_CANDIDATO|TRACE_REASONS_CANDIDATO|Res_Tit_Unificada|FAMILIA POR NOMBRE_TITULO|NIVEL DE SALIDA', str(worktree)]
search = subprocess.run(command, capture_output=True, text=True, timeout=45)
assert search.returncode in [0, 1]
result = {'fecha': '2026-10-26', 'proceso': 'Titulación/conciliación de mallas IPSS 162',
          'fuentes': sources, 'busqueda_codigo': {'ruta': str(worktree), 'comando': command,
          'exit_code': search.returncode, 'archivos_coincidentes': search.stdout.splitlines(), 'errores': search.stderr},
          'pendientes_e': ['Fórmula/pesos del puntaje 2130 no localizados; VLOOKUP no es ese algoritmo.',
                          'Atribución formal título–CODCLI–plan de los cuatro casos no demostrada.',
                          'CUADRO HOMOLOGACIÓN vincula programas con SIES, no ramos ni autorizaciones de convalidación.',
                          'No se obtuvieron nuevas equivalencias para los 21 pares carrera–plan–ramo.'],
          'limites': 'Inspección de fuentes y fórmulas, sin ejecutar Excel ni cambiar fuentes, diccionario, bases o denominadores; sin commit/push.'}
payload = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
# Garantía específica contra divulgación de las identidades leídas.
for c in cases:
    assert c['RUT'] not in payload and c['RUT'].split('-')[0] not in payload
    for t in c['trazas_calculadas']:
        assert t['codcli_candidato'] not in payload
(OUT / 'RESULTADO_AGREGADO.json').write_text(payload)
print(json.dumps([{'sha256': x['sha256'], 'validacion_traza': x['validacion_traza'],
                  'busqueda_auxiliar': x['inspeccion_busqueda_auxiliar_no_recalculo_excel']} for x in sources], ensure_ascii=False))
