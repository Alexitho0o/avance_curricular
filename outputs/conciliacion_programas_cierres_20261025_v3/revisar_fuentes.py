import json,re,hashlib,collections,datetime,csv
from pathlib import Path
from openpyxl import load_workbook
R=Path('/Users/alexi/Documents/GitHub/avance_curricular'); T=Path(__file__).parent
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def text(v):return '' if v is None else str(v).strip()
p=json.loads((T/'entrada_agregada.json').read_text()); master=json.loads((R/'data/versiones_historicas_sies_umas/fase_5_20261015/maestro.json').read_text())
slots=Path('/Users/alexi/.local/share/ipss/trazabilidad_integral_fase3_20261025/VISTA_T1_T2_T3.json')
targets={r[0]:r for r in p['foco']}; contexts=[c for s in json.loads(slots.read_text()) for c in s['contextos'] if c['plan_destino'] in targets and not c['codigo_destino']]
pm=Path('/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_12166.xlsx'); wanted={f['fila'] for c in contexts for f in c['fuentes_promedios']}
w=load_workbook(pm,read_only=True,data_only=True); s=w['Hoja1']; rows=iter(s.values);h=next(rows);found={i:dict(zip(h,r)) for i,r in enumerate(rows,2) if i in wanted};w.close()
periods=collections.defaultdict(lambda: {'n':0,'valid':0,'rows':set(),'states':set()})
for c in contexts:
 key=(c['plan_destino'],c['anio_observado'],c['periodo_nativo']); out=periods[key];out['n']+=1
 refs=[(f['fila'],found[f['fila']]) for f in c['fuentes_promedios']]
 valid=[(rn,d) for rn,d in refs if text(d['CODCLI'])==c['CODCLI'] and text(d['PLAN_DE_ESTUDIO'])==c['plan_destino'] and text(d['CODCARR'])==c['carrera_umas'] and text(d['ANO'])==str(c['anio_observado']) and text(d['PERIODO'])==str(c['periodo_nativo']) and text(d['JORNADA'])=='O']
 candidates=[a for a in master['oferta_anual'] if a['anio']==c['anio_observado'] and a['codigo_unico']==targets[c['plan_destino']][2]]
 # Program/period correspondence only: no documentary attribution or effective migration date asserted.
 if valid and candidates and str(c['sede_contexto'])=='2' and str(c['jornada_contexto'])=='4':out['valid']+=1
 for rn,d in valid:out['rows'].add(rn);out['states'].add(text(d['DESCRIPCION_ESTADO']) or 'Sin estado ramo')
temporal=[]
for (plan,y,per),d in sorted(periods.items()):temporal.append([plan,y,per,targets[plan][2],d['n'],d['valid'],'B/C: correspondencia contextual corroborada' if d['n']==d['valid'] else 'E: evidencia incompleta','PROMEDIOS Hoja1 filas '+','.join(map(str,sorted(d['rows']))),'; '.join(sorted(d['states'])),'E: no acredita envío SIES T2/T3 ni fecha de migración; no extrapolar'])
sources=[[str(pm),sha(pm)],[str(slots),sha(slots)]]
hs=json.loads((R/'outputs/auditoria_embudo_2317_20261022/HASHES_FUENTES.json').read_text());ac=collections.Counter();actotal={};acrefs={};tit=collections.Counter();titsets=collections.defaultdict(set);titrefs={};years=set();audit=[]
for path in hs:
 if not (path.endswith('.xlsx') and ('A_CURRICULAR_' in path or ('titulacio' in path and '/2026/' in path))):continue
 w=load_workbook(path,read_only=True,data_only=True);sources.append([path,sha(path)])
 if 'A_CURRICULAR_' in path:
  s=w['Matrícula_Informada']; rows=list(s.values);header=rows[6];codeidx=next(i for i,x in enumerate(header) if 'Código de la carrera asociado' in text(x)); desc=' '.join(text(v) for r in rows[:7] for v in r)
  yr=re.search(r'Durante el año académico (\d{4})',desc)
  if not yr:audit.append([path,'E: año académico no localizado']);continue
  y=int(yr[1]);n=0
  for r in rows[7:]:
   code=text(r[codeidx]);
   if re.fullmatch(r'I162S\d+C\d+J\d+V\d+',code):ac[(y,code)]+=1;n+=1
  actotal[y]=n;acrefs[y]=Path(path).name+' / Matrícula_Informada, encabezado fila 7';audit.append([path,f'B: referencia {y}; {n} filas con código completo; no comprobante de aceptación'])
 else:
  s=w['Res_Tit_Carga'];rows=iter(s.values);head=next(rows);n=missing=0
  for r in rows:
   d=dict(zip(head,r))
   if not d.get('TIPO_DOCUMENTO'):continue
   n+=1;code=text(d.get('COD_SIES_OBT_TIT'));date=d.get('FECHA_OBT_TIT_GRA'); y=date.year if isinstance(date,(datetime.date,datetime.datetime)) else int(str(date)[:4]) if re.match(r'^20\d\d',str(date)) else None
   if y and re.fullmatch(r'I162S\d+C\d+J\d+V\d+',code):
    titsets[(y,code)].add((text(d.get('TIPO_DOCUMENTO')),text(d.get('NUM_DOCUMENTO')),text(d.get('DV'))));tit[(y,code)]=len(titsets[(y,code)]);years.add(y);titrefs[y]=Path(path).name+' / Res_Tit_Carga; FECHA_OBT_TIT_GRA; personas distintas por código/año'
   else:missing+=1
  audit.append([path,f'B: {n} registros de titulación; {missing} sin código/fecha utilizable. Año de título separado de carga 2026. No prueba aceptación.'])
 w.close()
# Historical title extract: explicit year, no merging of overlapping source years.
titlepath=next(Path(x) for x in hs if 'titulacio' in x and x.endswith('.xlsx'))
capture=next(titlepath.parent.glob('16683*.xlsx'))
w=load_workbook(capture,read_only=True,data_only=True);s=next(s for s in w if s.title.startswith('5319'));rows=iter(s.values);head=next(rows);histtit=collections.defaultdict(set)
for r in rows:
 d=dict(zip(head,r));code=text(d.get('COD_UNICO_OBT_TIT'));y=text(d.get('AÑO_TITULACION'))
 if y.isdigit() and re.fullmatch(r'I162S\d+C\d+J\d+V\d+',code):histtit[(int(y),code)].add((text(d.get('TIPO_DOCUMENTO')),text(d.get('NUM_DOCUMENTO')),text(d.get('DV'))))
for (y,code),ids in histtit.items():
 if y not in years:tit[(y,code)]=len(ids);titrefs[y]=capture.name+' / '+s.title+'; AÑO_TITULACION; personas distintas por código/año'
years.update(y for y,c in histtit);sources.append([str(capture),sha(capture)]);audit.append([str(capture),'B/C: histórico 5319, personas distintas por código/año; no sumar con carga 2026 en años solapados']);w.close()
offers=collections.defaultdict(list)
for a in master['oferta_anual']:offers[a['codigo_unico']].append(a)
mf=json.loads((R/'outputs/trazabilidad_integral_umas_sies_20261025/fase2/conciliacion/MANIFEST.json').read_text());csvpath=next(x for x in mf['entradas_sha256'] if x.endswith('10_07_2026.csv'));assert sha(csvpath)==mf['entradas_sha256'][csvpath]
enrol={};names={}
with open(csvpath,encoding='latin-1') as f:
 for row in csv.DictReader(f,delimiter=';'):
  if row['CÓDIGO DE INSTITUCIÓN']!='162':continue
  key=(int(row['AÑO'].replace('MAT_','')),row['CÓDIGO CARRERA']);assert key not in enrol;enrol[key]=int(row['TOTAL MATRÍCULA'].replace('.',''));names[key[1]]=row['NOMBRE CARRERA']
sources.append([csvpath,sha(csvpath)])
annual={(a[0],a[1]):a for a in p['offers']};codes=sorted(set(offers)|{a[1] for a in p['offers']}|{c for y,c in tit}|{c for y,c in ac}|set(names))
detail=[]
for code in codes:
 hist=offers.get(code,[]); ys=sorted(a['anio'] for a in hist);first=min(ys) if ys else None;last=max(ys) if ys else None
 for y in range(2007,2027):
  a=annual.get((y,code));of=next((r for r in hist if r['anio']==y),None)
  status=of['vigencia'] if of else ('Sin oferta observada después de última aparición; cierre no acreditado' if last and y>last else 'Sin fila de oferta en cobertura revisada')
  name=of['nombre_carrera'] if of else hist[-1]['nombre_carrera'] if hist else names.get(code,'Nombre pendiente')
  detail.append([y,code,name,status,first,last,'Pendiente: no hay resolución/fecha de cierre acreditada',enrol.get((y,code),'Sin fila CSV'),a[13] if a and a[13] is not None else 'Sin dato',a[14] if a and a[14] is not None else 'Sin dato',None,tit.get((y,code),0) if y in years else 'Sin cobertura',ac.get((y,code),0) if y in actotal else 'Sin cobertura','No acreditado en esta revisión',titrefs.get(y,'Sin fuente anual revisada'),acrefs.get(y,'Sin fuente anual revisada'),'B/C; E cierre y aceptación'])
result={'temporal':temporal,'detail':detail,'sources':sources,'audit':audit,'qa':{'contextos':len(contexts),'corroborados':sum(r[5] for r in temporal),'filas_pm_verificadas':len(wanted),'codigos':len(codes),'filas_historico':len(detail),'ac_por_anio':actotal,'titulaciones_por_anio':{str(y):sum(n for (yy,c),n in tit.items() if yy==y) for y in years}}}
(T/'evidencia_temporal_historica.json').write_text(json.dumps(result,ensure_ascii=False));print(json.dumps(result['qa'],ensure_ascii=False))
