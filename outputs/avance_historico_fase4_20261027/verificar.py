"""QA independiente de salida: valores, fórmulas, RAW, hashes y alcance."""
import json, math, hashlib, subprocess, zipfile, re
from pathlib import Path
from collections import defaultdict
from xml.etree import ElementTree as ET
import openpyxl

O=Path(__file__).parent; R=O.parents[1]
P=Path.home()/'.local/share/ipss/avance_historico_fase4_20261027'
F=O/'AVANCE_ACADEMICO_HISTORICO_IPSS_20261027_V1.xlsx'
d=json.loads((P/'inputs.json').read_text()); base=json.loads((O/'BASELINE.json').read_text())
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
checks=[]
def ck(n,ok,detail=None):
    checks.append({'control':n,'pass':bool(ok),**({'detalle':detail} if detail is not None else {})})
def eq(a,b):
    return math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-10) if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
def ratio(n,v):return n/v if v else 'Pendiente'
w=openpyxl.load_workbook(F,read_only=True,data_only=True)
ck('15 hojas',len(w.sheetnames)==15)
rows=list(w['07_QA_matematico'].iter_rows(min_row=6,max_row=17,min_col=1,max_col=5,values_only=True))
for r in rows:ck('Matemática: '+r[0],eq(r[1],r[2]) and r[3]==0 and r[4]=='Coincide')
aggregate=defaultdict(lambda:[0]*6)
for r in d['annual']:
    x=aggregate[r[1]+'|'+r[2]];x[0]+=1
    for j,k in enumerate([6,7,9,10],1):x[j]+=r[k]
    x[5]+=next(p[4] for p in d['plans'] if p[0]==r[1]+'|'+r[2])
bad=0
for r in w['02_Anual_SIES_2025'].iter_rows(min_row=26,max_row=68,max_col=14,values_only=True):
    e=aggregate[r[0]]
    expected=[e[0],e[1],e[2],ratio(e[2],e[1]),e[3],e[4],e[5],ratio(e[3],e[5]),ratio(e[4],e[5])]
    got=[r[i] for i in [4,6,7,8,9,10,11,12,13]]
    bad+=sum(not eq(a,b) for a,b in zip(expected,got))
ck('43 desgloses anuales × 9 resultados',bad==0,{'diferencias':bad})
plan_lookup={x[0]:x for x in d['plans']};bad=0
for r,a in zip(w['08_Base_anual'].iter_rows(min_row=6,max_row=2376,max_col=20,values_only=True),d['annual']):
    p=plan_lookup[a[1]+'|'+a[2]]
    bad+=r[15]!=p[4] or r[16]!=p[3]
ck('2371 búsquedas de unidades y nombre exactos',bad==0,{'diferencias':bad})
bad=0
for a,b in zip(w['03_Periodos_UMAS'].iter_rows(min_row=6,max_row=5+len(d['periods']),max_col=13,values_only=True),d['periods']):
    expect=[b[6],b[9],b[10],b[11]+b[12],b[13],ratio(b[9],b[9]+b[10])]
    bad+=sum(not eq(x,y) for x,y in zip(a[6:12],expect))
ck('2758 contextos actividad × 6 resultados',bad==0,{'diferencias':bad})
bad=0
for a,b in zip(w['04_Mallas'].iter_rows(min_row=26,max_row=25+len(d['mallas']),max_col=16,values_only=True),d['mallas']):
    bad+=not eq(a[13],ratio(b[10],b[7]))
    bad+=a[14]!='Pendiente'
    for i in [4,6,7,8,9,10,11,12]:bad+=not eq(a[i],b[i])
ck('565 ratios de malla y estado pendiente',bad==0,{'diferencias':bad})
selected=[r for r in d['mallas'] if r[0]=='prom12166']
expect=[sum(r[4] for r in selected if r[5]=='LISTADO_EXACTO_DISPONIBLE_VIGENCIA_PENDIENTE')]
expect += [sum(r[i] or 0 for r in selected) for i in [7,8,10]]
got=[r[0] for r in w['04_Mallas'].iter_rows(min_row=6,max_row=9,min_col=2,max_col=2,values_only=True)]
ck('Resumen de mallas responde a fuente seleccionada',all(eq(x,y) for x,y in zip(expect,got)),{'esperado':expect,'observado':got})
bad=0
for r in w['01_Resumen'].iter_rows(min_row=31,max_row=30+len(d['careers']),max_col=9,values_only=True):
    selected=[x for x in d['periods'] if x[0]=='prom12166' and x[1]=='2026' and x[2]=='2' and x[3]==r[0]]
    e=[sum(x[i] for x in selected) for i in [6,9,10]]+[sum(x[11]+x[12] for x in selected),sum(x[13] for x in selected)]
    bad+=sum(not eq(x,y) for x,y in zip(e,r[2:7]))
ck('Desglose ejecutivo por carrera',bad==0,{'carreras':len(d['careers']),'diferencias':bad})
# RAW: identidad, orden y valores completos; no se escriben identificadores en el QA.
for sheet,label in [('13_RAW_7804','prom7804'),('14_RAW_12166','prom12166')]:
    errors=0; n=0
    with Path(d['raw'][label]['ruta']).open() as f:
        for n,(a,line) in enumerate(zip(w[sheet].iter_rows(min_row=6,max_col=17,values_only=True),f),1):
            b=json.loads(line)
            errors+=any(('' if x is None else x)!=('' if y is None else y) for x,y in zip(a,b))
    ck('RAW íntegro '+label,n==d['raw'][label]['filas'] and errors==0 and w[sheet].max_row==n+5,{'filas':n,'filas_distintas':errors})
bad=0;n=0
for n,(a,b) in enumerate(zip(w['12_Listado_mallas'].iter_rows(min_row=6,max_col=11,values_only=True),d['catalog']),1):
    bad+=any(('' if x is None else x)!=('' if y is None else y) for x,y in zip(a,b))
ck('Listado de mallas íntegro',n==len(d['catalog']) and bad==0 and w['12_Listado_mallas'].max_row==n+5,{'filas':n,'filas_distintas':bad})
w.close()
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
with zipfile.ZipFile(F) as z:
    ck('ZIP íntegro',z.testzip() is None)
    fc=cache=errors=blank_expected=0
    for n in z.namelist():
        if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):
            # Streaming parse evita materializar las dos bases de gran tamaño.
            with z.open(n) as f:
                for _,el in ET.iterparse(f,events=['end']):
                    if el.tag=='{'+NS['s']+'}c':
                        if el.find('s:f',NS) is not None:
                            fc+=1;cache+=el.find('s:v',NS) is not None
                            errors+=el.attrib.get('t')=='e'
                            if el.find('s:v',NS) is None and el.attrib.get('t')=='str':
                                formula=el.find('s:f',NS).text or ''
                                m=re.match(r"IF\(ISBLANK\('11_Base_mallas'!([A-N])(\d+)\)",formula)
                                if m and d['mallas'][int(m[2])-6][ord(m[1])-65] is None:blank_expected+=1
                        el.clear()
    ck('Fórmulas evaluadas y sin error',fc>0 and fc==cache+blank_expected and errors==0,{'formulas':fc,'con_cache_no_vacio':cache,'blancos_esperados_verificados':blank_expected,'errores':errors})
    charts=[n for n in z.namelist() if re.fullmatch(r'xl/(?:drawings/)?charts/chart\d+\.xml',n)]
    ck('3 gráficos editables',len(charts)==3)
    ck('Gráficos con referencia a celdas',all(ET.fromstring(z.read(n)).find('.//c:f',NS) is not None for n in charts))
    tables=[n for n in z.namelist() if n.startswith('xl/tables/table') and n.endswith('.xml')]
    ck('Tablas con encabezados únicos',all(len([x.attrib['name'] for x in ET.fromstring(z.read(n)).find('s:tableColumns',NS)])==len(set(x.attrib['name'] for x in ET.fromstring(z.read(n)).find('s:tableColumns',NS))) for n in tables))
    ck('Tabla de mallas cubre 11438 filas',any(ET.fromstring(z.read(n)).attrib.get('name')=='ListadoMallas' and ET.fromstring(z.read(n)).attrib.get('ref')=='A5:K11443' for n in tables))
    ck('Sin vínculos externos','xl/externalLinks/' not in '\n'.join(z.namelist()))
motor=json.loads((O/'QA_MOTOR.json').read_text()); fil=[x for x in motor if x['control']=='filtro']
ck('6 pruebas de selección con recálculo',len(fil)==6 and all(x['pass'] for x in fil))
for k,v in base['sources'].items():ck('Fuente preservada: '+k,sha(v['ruta'])==v['sha256'])
ck('V24 preservada',sha(base['reference']['ruta'])==base['reference']['sha256'])
for k,v in base['protected'].items():ck('Cambio previo preservado: '+k,sha(R/k)==v)
head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=R).decode().strip()
branch=subprocess.check_output(['git','branch','--show-current'],cwd=R).decode().strip()
ck('HEAD y rama sin cambios',head==base['git_head'] and branch==base['git_branch'])
ck('Excel excluido de Git',subprocess.run(['git','check-ignore','-q',str(F)],cwd=R).returncode==0)
result={'clasificacion':'C','estado':'PASS_CON_ALCANCE' if all(x['pass'] for x in checks) else 'FAIL','controles':checks,'fuentes_y_denominadores_preservados':True,'excel_nativo':'NO_VERIFICADO','sha256':sha(F),'head':head,'rama':branch,'limites':['No certifica vigencia histórica de mallas ni cobertura total de extractos.','No aplica las 644 atribuciones candidatas.','Recálculo con Artifact Tool y contraste independiente; prueba de apertura en Excel nativo separada.']}
(O/'QA_FINAL.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'estado':result['estado'],'controles':len(checks),'fallidos':[x for x in checks if not x['pass']],'sha256':result['sha256']},ensure_ascii=False))
raise SystemExit(0 if result['estado']!='FAIL' else 1)
