import openpyxl,csv,json,hashlib
from pathlib import Path
from collections import Counter
O=Path(__file__).parent;B=Path('data/reglas_negocio/20261022_puente_versiones/validacion_local_v1');src=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
w=openpyxl.load_workbook(src,read_only=True,data_only=True)
blocks=[]
# Selección explícita de columnas no personales. Se cuentan filas, no personas.
specs={9:('Cohorte de ingreso',[1,2,4,5]),11:('ANIO_INFORMADO',[0,11,12,14,15,32]),20:('AÑO',[1,2,9,10,11]),24:('CODCLI',[3,5,6,9,10]),25:('CODCLI',[3,5,6,9,10])}
for index,(marker,cols) in specs.items():
 s=w.worksheets[index];header=None;count=Counter()
 for row in s.iter_rows(values_only=True):
  if marker in row:header=list(row);continue
  if header is not None and any(row[c] is not None for c in cols) and isinstance(row[cols[0]],(str,int,float)):
   count[tuple(row[c] for c in cols)]+=1
 assert header
 rows=[[s.title+' · agregado provisional'],['2026-10-22 · detalle individual excluido; original preservado'],['C · Cantidad de registros de la fuente; no se interpreta como personas únicas'],[str(header[c]) for c in cols]+['Registros']]+[list(k)+[v] for k,v in sorted(count.items(),key=lambda kv:str(kv[0]))]
 blocks.append({'name':s.title,'replace_index':index+1,'rows':rows})
summary=[['Actualización provisional · puente UMAS–SIES'],['2026-10-22 · PENDIENTE_APROBACION_GOBERNANZA'],['Cobertura del puente y retención son métricas diferentes. No se demuestra reducción de brecha.'],['Indicador','Antes n','Antes d','Antes %','Ahora n','Ahora d','Ahora %','Cambio pp'],['Propuestas internas previas (no aprobadas)',39,63,None,39,63,None,None],['Pendientes múltiples/baja cobertura',11,63,None,11,63,None,None],['Sin match en códigos candidatos',13,63,None,13,63,None,None],['Brecha de propuestas (no equivalencias)',24,63,None,24,63,None,None],[],['Mejora demostrada','9 de 13 sin match tienen desfase temporal documentado; no implica equivalencia.'],['QA','Dos implementaciones coinciden; cero diferencia en totales MU.'],['Discrepancias','IECIRE S2/J4/V2: 23→22; IECIREOL: 243→241. Causa pendiente.'],['Retención histórica','Numeradores, denominadores y tasas fuente preservados. No se recalcularon por el puente.'],['Uso de esta copia','Históricos afectados por excluir datos individuales quedan como foto de valores; filtros no recalculan esos resultados.'],['Datos SIES','Tablas oficiales y comparativos conservados; propuestas internas no son reglas oficiales.'],['Pendientes','Tipo/duración UMAS, temporalidad, dos diferencias de conteo y conciliación semántica de retención.']]
blocks.append({'name':'29_Actualizacion_Provisional','rows':summary,'formulas':{f'{col}{r}':f'={num}{r}/{den}{r}' for r in range(5,9) for col,num,den in [('D','B','C'),('G','E','F')]} });blocks[-1]['formulas'].update({f'H{r}':f'=(G{r}-D{r})*100' for r in range(5,9)})
a=list(csv.DictReader(open(B/'TABLA_63_DIMENSIONES.csv',encoding='utf-8-sig')))
cols=['carrera_interna','sede','jornada','planes','estado_anterior','personas_coincidentes_unicas','distribucion_versiones','proporcion_incidencias','estado','pendiente']
blocks.append({'name':'30_Puente_63_Dimensiones','rows':[['Puente interno · 63 dimensiones'],['C/E · propuestas y evidencia; no equivalencias aprobadas'],['PROMEDIOS 2019–2026; MU 2022–2026; catálogo 2027'],cols]+[[float(r[c]) if c=='proporcion_incidencias' and r[c] else int(r[c]) if c=='personas_coincidentes_unicas' else r[c] for c in cols] for r in a]})
rows=[['Retención histórica · antes y ahora'],['B/C · copia del libro seleccionado. Comparación independiente; no cambia tasas históricas.'],['Se conserva tasa fuente incluso si difiere de numerador/denominador; discrepancia explícita.'],['Fuente / grupo','Año','Numerador antes','Denominador antes','Tasa fuente antes','Numerador ahora','Denominador ahora','Tasa fuente ahora','Brecha fuente pp','Tasa n/d','Desfase fuente vs n/d pp']]
for row in w['4_Resultados_Sector'].iter_rows(min_row=7,max_row=30,values_only=True):
 if isinstance(row[2],int) and isinstance(row[3],(int,float)) and isinstance(row[4],(int,float)):rows.append(['IPSS · '+str(row[1]),row[2],row[4],row[3],row[5],row[4],row[3],row[5]])
for row in w['4_Resultados_Sector'].iter_rows(min_row=7,max_row=9,values_only=True):rows.append(['SIES oficial · '+str(row[7]),2025,row[12],row[11],row[13],row[12],row[11],row[13]])
for rn in [7,8,9,10,11,12]:
 row=next(w['18_Foto_Intertrimestral'].iter_rows(min_row=rn,max_row=rn,values_only=True))
 if isinstance(row[6],(int,float)):rows.append(['IPSS Foto T3 · '+str(row[2]),row[1],row[6],row[3],row[7],row[6],row[3],row[7]])
formula={}
for r in range(5,len(rows)+1):formula.update({f'I{r}':f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(H{r})),(H{r}-E{r})*100,0)',f'J{r}':f'=IF(D{r}>0,C{r}/D{r},0)',f'K{r}':f'=IF(AND(D{r}>0,ISNUMBER(E{r})),(E{r}-J{r})*100,0)'})
blocks.append({'name':'31_Comparacion_Historica','rows':rows,'formulas':formula})
blocks.append({'name':'32_Fuentes_y_Exclusiones','rows':[['Fuentes y alcance de la copia provisional'],['Estado','PENDIENTE_APROBACION_GOBERNANZA'],['Base elegida',str(src)],['SHA-256 base',hashlib.sha256(src.read_bytes()).hexdigest()],['Fecha','2026-10-22'],['Criterio','Último libro AUDITADO_ENTREGA elegido por el usuario, con 13 gráficos y comparativos SIES.'],['Exclusión','Hojas 9,11,20,24,25: registros individuales sustituidos por agregados en la misma posición.'],['Conservación','Las demás hojas y los 13 gráficos se preservan. Fórmulas dependientes del detalle excluido se fijan al valor fuente.'],['Advertencia','Esta copia agregada es una foto descriptiva. No usar controles de selección para recalcular los históricos fijados.'],['Evidencia puente',str(B.resolve())],['Clasificaciones','A oficial; B observado; C cálculo; D decisión interna; E pendiente. No se crean nuevas reglas A.'],['Brecha puente','24/63 = 38,0952% antes y ahora; reducción demostrada: 0 pp.'],['Pendientes','Discrepancias 23→22 y 243→241 sin causa probada. No afectan tasas históricas.'],['Privacidad','Sin identificadores personales. RAW de datos agregados y SIES se conserva.']]})
(O/'payload.json').write_text(json.dumps(blocks,ensure_ascii=False,default=str))
print('bloques',len(blocks),'comparaciones',len(rows)-4)
