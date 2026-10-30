"""H003: fuentes fijadas -> modelo agregado -> Excel. No recalcula H001/H002."""
from pathlib import Path
import csv, hashlib, json, re, subprocess, sys, copy
import openpyxl
from openpyxl.formula.translate import Translator

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
RUNTIME = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
BOOK = 'ENTREGA_RETENCION_IPSS_20261029.xlsx'
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p, d): Path(p).write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
def verify_sources():
    sources=json.loads((OUT/'FUENTES_FIJADAS.json').read_text())
    for s in sources:
        if sha(s['ruta']) != s['sha256']: raise ValueError('Hash distinto: '+s['id'])
    expected={'V24':'bfeba676','V24_COPIA_RECALCULADA':'1e28d595','SIES_TASAS':'03897a16'}
    for s in sources:
        if s['id'] in expected: assert s['sha256'].startswith(expected[s['id']])
    return {s['id']:s for s in sources}

def presentar_h003b(sheets, controls):
    """Reubica bloques; conserva cada dato y la lógica de las fórmulas.

    Único cambio de fuente H003b: este archivo. El autor JS y el verificador
    H003 se reutilizan sin modificarlos. Las direcciones cambian al mover una
    tabla, con traducción relativa y una comprobación por cada celda original.
    """
    original=copy.deepcopy(controls)
    moves={}
    def move(s,b,row,col=0,order=None):
        oldrow=b['row']; order=order or list(range(len(b['rows'])))
        positions={old:new for new,old in enumerate(order)}
        for oldi,newi in positions.items():
            for j in range(len(b['headers'])):
                old=f'{openpyxl.utils.get_column_letter(j+1)}{oldrow+1+oldi}'
                new=f'{openpyxl.utils.get_column_letter(j+col+1)}{row+1+newi}'
                moves[(s['name'],old)]=new
        b['formulas']={moves[(s['name'],a)]:Translator(f,origin=a).translate_formula(moves[(s['name'],a)]) for a,f in b['formulas'].items()}
        b['formats']={openpyxl.utils.get_column_letter(openpyxl.utils.column_index_from_string(c)+col):v for c,v in b['formats'].items()}
        b['rows']=[b['rows'][i] for i in order]; b['row']=row; b['startCol']=col

    s=sheets[1]
    s['widths']=[12,16,16,17,17,19,58,20,4,34,15,16,17,15,16,17,20]
    s['contentEnd']='H'
    s['printAreas']=['A1:H29','J5:Q14']
    move(s,s['blocks'][1],6,9)
    for n in s['notes']:
        if n['row']==21: n.update(row=12,start='J',end='Q')
    s['renderRanges']=['A1:H29','J5:Q14']

    s=sheets[3]
    s['widths'][0]=52
    summary,detail=s['blocks'][:2]
    move(s,summary,6)
    groups={'CONTINUA':0,'VERSION_MULTIPLE_RESUELTA':0,'APROBADA_D_2027':0,
            'NUEVO_2027_SIN_PLAN_E':1,'HISTORICO_SIN_MATRICULA_RECIENTE':1,'FUERA_ALCANCE':2}
    order=sorted(range(len(detail['rows'])),key=lambda i:(groups[detail['rows'][i][3]],detail['rows'][i][0]))
    move(s,detail,20,order=order)
    detail['title']='Tabla completa 2027 · resueltos, pendientes y fuera de alcance'
    for n in s['notes']:
        if n['row']==183: n['row']=17
    s['freeze']=20
    s['renderRanges']=['A1:H27','A20:H39']

    for x in controls:
        old=x['celda']; new=moves.get((x['hoja'],old),old)
        x['celda']=new
        if 'formula' in x: x['formula']=Translator(x['formula'],origin=old).translate_formula(new)
    assert len(controls)==len(original)==1956
    for before,after in zip(original,controls):
        assert all(before.get(k)==after.get(k) for k in ['hoja','esperado','fuente','tipo'])
        if 'formula' in before:
            assert Translator(after['formula'],origin=after['celda']).translate_formula(before['celda'])==before['formula']
    dump(OUT/'QA_PRESERVACION_H003B.json',{'estado':'PASS','controles':1956,'datos_fuentes_clases_sin_cambios':True,
         'formulas_misma_logica':True,'celdas_reubicadas':sum(a['celda']!=b['celda'] for a,b in zip(original,controls)),
         'detalle':[{'hoja':a['hoja'],'antes':a['celda'],'despues':b['celda'],'coincide':True} for a,b in zip(original,controls)]})

def compactar_h003c(sheets, controls):
    """Una fila de separación entre secciones; notas finales contiguas.

    Solo cambia la posición de bloques. Conserva textos y la lógica relativa
    de las 41 fórmulas, con correspondencia de las 1.956 celdas controladas.
    """
    before=copy.deepcopy(controls); moves={}; changes=[]
    for s in sheets:
        if s['name']=='1_Retencion_1er_anio':
            bottom=max(b['row']+len(b['rows']) for b in s['blocks'])
            for i,n in enumerate(n for n in s['notes'] if n['row']>=5):
                changes.append({'hoja':s['name'],'nota_antes':n['row'],'nota_despues':bottom+2+i})
                n.update(row=bottom+2+i,start='A',end='Q',span=1)
            s['printAreas']=['A1:Q14'];s['renderRanges']=['A1:Q14']
            continue
        events=[(b['row']-1,'tabla',b) for b in s['blocks']]
        events += [(n['row'],'nota',n) for n in s['notes'] if n['row']>=5]
        last=3; previous=None
        for oldstart,kind,obj in sorted(events,key=lambda x:x[0]):
            start=5 if previous is None else last+(1 if kind==previous=='nota' else 2)
            if kind=='nota':
                changes.append({'hoja':s['name'],'nota_antes':obj['row'],'nota_despues':start})
                obj.update(row=start,span=1);last=start
            else:
                oldrow=obj['row'];newrow=start+1;shift=newrow-oldrow
                for i in range(len(obj['rows'])):
                    for j in range(len(obj['headers'])):
                        col=openpyxl.utils.get_column_letter(j+1+obj.get('startCol',0))
                        moves[(s['name'],f'{col}{oldrow+1+i}')]=f'{col}{newrow+1+i}'
                obj['row']=newrow;last=newrow+len(obj['rows'])
                changes.append({'hoja':s['name'],'tabla':obj['title'],'fila_antes':oldrow,'fila_despues':newrow})
            previous=kind
    def remap_formula(formula,sheet,mapping):
        # Al mover bloques independientes, una referencia a otro bloque conserva
        # su precedente real; no debe trasladarse como una copia de celda.
        from openpyxl.formula.tokenizer import Tokenizer
        parts=[]
        for token in Tokenizer(formula).items:
            value=token.value
            if token.subtype=='RANGE':
                assert '!' not in value, 'Revisar referencia entre hojas'
                value=re.sub(r'\b[A-Z]{1,3}[0-9]+\b',lambda m:mapping.get((sheet,m[0]),m[0]),value)
            parts.append(value)
        return '='+''.join(parts)
    for s in sheets:
        for b in s['blocks']:
            b['formulas']={moves.get((s['name'],a),a):remap_formula(f,s['name'],moves) for a,f in b['formulas'].items()}
    for x in controls:
        old=x['celda'];new=moves.get((x['hoja'],old),old);x['celda']=new
        if 'formula' in x:x['formula']=remap_formula(x['formula'],x['hoja'],moves)
    assert len(controls)==len(before)==1956
    for a,b in zip(before,controls):
        assert all(a.get(k)==b.get(k) for k in ['hoja','esperado','fuente','tipo'])
        if 'formula' in a:assert remap_formula(b['formula'],b['hoja'],{(s,new):old for (s,old),new in moves.items()})==a['formula']
    dump(OUT/'QA_COMPACTACION_H003C.json',{'estado':'PASS','controles':1956,'diferencias':0,'cambios_presentacion':changes,
         'detalle':[{'hoja':a['hoja'],'antes':a['celda'],'despues':b['celda']} for a,b in zip(before,controls)]})

def autor_h003b():
    """Extiende en memoria el autor existente; no escribe otro script JS."""
    js=(OUT/'construir_excel.mjs').read_text()
    def change(old,new):
        nonlocal js
        assert old in js,old
        js=js.replace(old,new)
    change("const out=path.dirname(fileURLToPath(import.meta.url));",'const out='+json.dumps(str(OUT))+';')
    change("sh.mergeCells(`A2:${end}2`)","sh.mergeCells(`A2:${spec.contentEnd||end}2`)")
    change("sh.getRange(`A2:${end}2`)","sh.getRange(`A2:${spec.contentEnd||end}2`)")
    change("`A${note.row}:${end}${note.row+1}`","`${note.start||'A'}${note.row}:${note.end||spec.contentEnd||end}${note.row+1}`")
    change('${note.row+1}', '${note.row+(note.span||2)-1}')
    change("rg.format.font={name:model.font,size:11,color:colors.azul_institucional};", """rg.format.font={name:model.font,size:11,color:colors.azul_institucional};
    if(note.span===1){
      const first=(note.start||'A').charCodeAt(0)-65;
      const final=(note.end||spec.contentEnd||end).charCodeAt(0)-65;
      const capacity=spec.widths.slice(first,final+1).reduce((a,b)=>a+b,0)*0.80-4;
      rg.format.rowHeight=Math.max(23,Math.ceil(note.text.length/capacity)*16+8);
    }""")
    change("const b=spec.blocks[k], be=col(b.headers.length-1), bottom=b.row+b.rows.length;",
           "const b=spec.blocks[k], start=b.startCol||0, bs=col(start), be=col(start+b.headers.length-1), bottom=b.row+b.rows.length;")
    change('`A${b.row-1}:${be}${b.row-1}`','`${bs}${b.row-1}:${be}${b.row-1}`')
    change('`A${b.row}:${be}${bottom}`','`${bs}${b.row}:${be}${bottom}`')
    change('`A${b.row}:${be}${b.row}`','`${bs}${b.row}:${be}${b.row}`')
    change('`A${b.row+1+i}:${be}${b.row+1+i}`','`${bs}${b.row+1+i}:${be}${b.row+1+i}`')
    change('spec.widths[j]*0.85','(spec.widths[start+j]-3)*0.80')
    change('sh.getCell(b.row+i,j)','sh.getCell(b.row+i,start+j)')
    change('for(let k=0;k<spec.blocks.length;k++){','const rowHeights=new Map();\n  for(let k=0;k<spec.blocks.length;k++){')
    change('row.format.rowHeight=Math.max(27,lines*15+10);',
           'const height=Math.max(rowHeights.get(b.row+1+i)||0,27,lines*15+10); rowHeights.set(b.row+1+i,height); row.format.rowHeight=height;')
    change('sh.freezePanes.freezeRows(spec.freeze);', """
  if(spec.name==='1_Retencion_1er_anio')sh.getRange('A7:A10').format.horizontalAlignment='center';
  if(spec.name==='3_Paridad_SIES_2027'){
    sh.getRange('A1:H1').format.rowHeight=10;
    sh.getRange('A3:H4').format.rowHeight=17;
    sh.getRange('A7:C15').format.rowHeight=22;
    sh.getRange('A6:C6').format.rowHeight=30;
    sh.getRange('A16:H16').format.rowHeight=9;
    sh.getRange('A17:H18').format.rowHeight=20;
  }
  sh.freezePanes.freezeRows(spec.freeze);""")
    change("const blob=await wb.render({sheetName:s.name,range:`A1:${s.lastCol}${Math.min(18,s.lastRow)}`,scale:1,format:'png'});\n  await fs.writeFile(path.join(out,'renders_artifact',s.name+'.png'),new Uint8Array(await blob.arrayBuffer()));",
           """const spec=model.sheets.find(x=>x.name===s.name);
  const ranges=spec.renderRanges||[`A1:${s.lastCol}${Math.min(18,s.lastRow)}`];
  for(let i=0;i<ranges.length;i++){
    const blob=await wb.render({sheetName:s.name,range:ranges[i],scale:1,format:'png'});
    await fs.writeFile(path.join(out,'renders_artifact',s.name+(i?'_'+(i+1):'')+'.png'),new Uint8Array(await blob.arrayBuffer()));
  }""")
    subprocess.run([str(RUNTIME/'node/bin/node'),'--input-type=module'],input=js,text=True,check=True,cwd=OUT)

def margen_texto_xlsx():
    """Sangría de presentación OOXML (sin API equivalente documentada).

    Conserva valores y fórmulas. Un punto de sangría evita que cifras y texto
    de columnas adyacentes se lean como una sola cadena.
    """
    from verificar_entrega import replace_zip,tag,ET
    import zipfile
    book=OUT/BOOK
    with zipfile.ZipFile(book) as z:
        root=ET.fromstring(z.read('xl/styles.xml'))
        for xf in root.find(tag('cellXfs')):
            alignment=xf.find(tag('alignment'))
            if alignment is None: alignment=ET.SubElement(xf,tag('alignment'))
            if alignment.get('horizontal') not in ['center','centerContinuous']:
                if not alignment.get('horizontal'):
                    alignment.set('horizontal','right' if xf.get('numFmtId','0')!='0' else 'left')
                alignment.set('indent','1');xf.set('applyAlignment','1')
    replace_zip(book,{'xl/styles.xml':ET.tostring(root,encoding='utf-8',xml_declaration=True)})

def verificar_h003b():
    """Reutiliza el QA intacto, con áreas de impresión por bloque independiente."""
    import verificar_entrega as qa
    import zipfile
    original=qa.print_setup
    def setup(book,layouts):
        original(book,layouts)
        model=json.loads((OUT/'MODELO_AGREGADO.json').read_text())
        with zipfile.ZipFile(book) as z:root=qa.ET.fromstring(z.read('xl/workbook.xml'))
        for dn in root.find(qa.tag('definedNames')):
            if dn.get('name')=='_xlnm.Print_Area':
                s=model['sheets'][int(dn.get('localSheetId'))]
                if 'printAreas' in s:dn.text=','.join("'"+s['name']+"'!"+a for a in s['printAreas'])
        qa.replace_zip(book,{'xl/workbook.xml':qa.ET.tostring(root,encoding='utf-8',xml_declaration=True)})
    qa.print_setup=setup
    qa.main()
    # CSV agregado versionado: fin de línea LF, sin alterar ningún campo.
    control_csv=OUT/'CONTROL_RESULTADOS.csv'
    data=control_csv.read_bytes()
    control_csv.write_bytes(data.replace(b'\r\n',b'\n'))
def main():
    src=verify_sources()
    def js(id): return json.loads(Path(src[id]['ruta']).read_text())
    cache=openpyxl.load_workbook(src['V24_COPIA_RECALCULADA']['ruta'],read_only=True,data_only=True)
    oficial=openpyxl.load_workbook(src['SIES_TASAS']['ruta'],read_only=True,data_only=True)
    sheets=[]; controls=[]
    def sheet(name,title,widths):
        s=dict(name=name,title=title,widths=widths,blocks=[],notes=[],freeze=6);sheets.append(s);return s
    def table(s,row,title,headers,rows,source,formats=None,formulas=None):
        b=dict(row=row,title=title,headers=headers,rows=rows,source=source,formats=formats or {},formulas=formulas or {})
        s['blocks'].append(b)
        for i,r in enumerate(rows):
            for j,v in enumerate(r):
                cell=f'{openpyxl.utils.get_column_letter(j+1)}{row+1+i}'
                if v is not None:
                    controls.append(dict(hoja=s['name'],celda=cell,esperado=v,fuente=source,tipo='valor'))
        return b
    def note(s,row,text): s['notes'].append(dict(row=row,text=text))
    def formula(b,col,index,text,expected,s):
        address=f'{col}{b["row"]+index+1}';b['formulas'][address]=text
        controls.append(dict(hoja=s['name'],celda=address,esperado=expected,fuente=b['source'],tipo='formula',formula=text))

    s=sheet('0_LEEME','IPSS · Retención y conciliaciones',[29,110]);s['freeze']=6
    table(s,6,'Alcance y lectura',['Tema','Alcance'],[
        ['Entrega','Preparada el 29-10-2026 para entrega el 31-10-2026. Solo agregados; sin bases individuales.'],
        ['Retención anual','Cohortes 2022–2025, observadas en el año siguiente. H001 interno; tasas publicadas SIES como referencia externa.'],
        ['Intertrimestral','Año académico 2026. Foto T3 al 22-10-2026, período en curso. Borrador interno; no tasa oficial SIES.'],
        ['Paridad de códigos','Catálogo 2026–2027, regla interna activa H002l. Oferta 2027 V06 preparada (C); no acredita aceptación PES.'],
        ['Cuadratura MU','Matrículas informadas 2022–2026. La brecha 2026 está cerrada documentalmente, con causa E.'],
        ['Unidades','Persona por cohorte en H001. Personas o matrículas según cada fila intertrimestral. Códigos y opciones en H002; no son estudiantes.'],
        ['Fórmulas','Tasas = numerador / denominador; diferencia = (tasa interna − tasa SIES) × 100. No promediar porcentajes ni sumar unidades distintas.'],
        ['Fuentes y commits','Rutas, SHA-256 y commits en 5_Fuentes. Los insumos se verifican antes de cada generación.'],
        ['Corte de fuentes','Se conservan los cortes de cada fuente, no una fecha de corte institucional común. Consulta el bloque correspondiente.'],
        ['Formato y marca','Paleta IPSS oficial. Arial sustituye a Montserrat/Gotham no instaladas. Porcentajes verificados en es-CL; el visor puede aplicar su configuración regional.'],
        ['Cambios protegidos','No se modifican fuentes, bases, denominadores, tasas ni reglas. No hay carga a SIES/PES ni commit/push.'],
    ],'H003; H001; INTER; H002_CIERRE; BRECHA')
    table(s,21,'Clases de respaldo',['Clase','Significado'],[['A','Regla oficial respaldada por manual, instructivo o comunicación oficial.'],['B','Dato observado en una fuente identificada; no constituye una regla.'],['C','Cálculo, implementación o inferencia técnica con alcance declarado.'],['D','Decisión interna aprobada; no sustituye una norma oficial.'],['E','Causa no acreditada, límite o pendiente expresamente conservado.']],'Instrucciones del proyecto')
    table(s,30,'Pendientes visibles',['Tema','Pendiente y forma de cierre'],[
        ['H001','Brechas 2022 y 2024 cerradas E: la investigación no acreditó su causa. No se ajustan tasas para igualar SIES.'],
        ['Intertrimestral','Paridad desde fuentes crudas, proxies y desgloses pendientes; cinco personas con matrícula duplicada no están en V24.'],
        ['Paridad 2027','34 opciones de revisión y 37 códigos nuevos; hay solapamiento. Resolver con evidencia y MU 2027; oferta aceptada antes de carga.'],
        ['MU 2026','Reabrir únicamente con fuente individual SIES de MU 2026, por ejemplo precarga AC 2027 cuando exista.'],
    ],'H001_CIERRE; INTER/cierre_hito_2c; REVISION/totales; BRECHA/reapertura')

    h=js('H001'); s=sheet('1_Retencion_1er_anio','Retención de 1er año · cohortes 2022–2025',[26,14,14,16,16,17,55,15])
    note(s,3,'Indicador interno alineado a la definición SIES; la tasa oficial es la publicada por SIES. Unidad: personas por cohorte.')
    tasas=[oficial['Retención 1er año x IES'].cell(59,c).value for c in range(18,22)]
    assert all(isinstance(x,(int,float)) and 0<x<1 for x in tasas)
    inst=h['grupos']['institucional']
    notes={2022:'E cerrada: +4,49 pp; causa no establecida tras revisar reglas y cobertura.',2023:'C: próxima a SIES; no idéntica (−0,08 pp).',2024:'E cerrada: −9,08 pp; 180 ausentes en RAW 2025, sin causa oficial demostrada.',2025:'C/E: próxima a SIES (−0,12 pp); 13 casos y 9 retenidos menos; causa individual E.'}
    b=table(s,6,'Comparación institucional',['Cohorte','Retenidos H001','Casos H001','Tasa H001','Tasa SIES publicada','Diferencia (pp)','Lectura','Clase'],[[x['cohorte'],x['retenidos'],x['casos'],None,tasas[i],None,notes[x['cohorte']],'C/D; B SIES; E'] for i,x in enumerate(inst)],'H001/grupos/institucional; SIES_TASAS/Retención 1er año x IES!R59:U59; H001_CIERRE/adendas',{'A':'year','B':'count','C':'count','D':'pct','E':'pct','F':'pp'})
    for i,x in enumerate(inst):
        r=7+i;formula(b,'D',i,f'=B{r}/C{r}',x['retenidos']/x['casos'],s);formula(b,'F',i,f'=(D{r}-E{r})*100',(x['tasa']-tasas[i])*100,s)
    # SIES agregado por sector: reutiliza mapeo aprobado; nunca suma Total general.
    mapping=js('SECTORES')['mapeo_por_cod_car']; totals={}
    for r in csv.DictReader(Path(src['SIES_GRANULAR']['ruta']).open(),delimiter='\t'):
        m=re.fullmatch(r'I162S\d+C(\d+)J\d+V\d+',r['mc_codigo_unico'])
        if not m: continue
        sec=mapping[m[1]][1]; t=totals.setdefault(sec,[0,0])
        for j,key in enumerate(['ret_1_retenido','N° de casos']): t[j]+=0 if r[key]=='-' else int(r[key])
    assert sum(v[0] for v in totals.values())==677 and sum(v[1] for v in totals.values())==1163
    sectors=[x for x in h['grupos']['sector'] if x['cohorte']==2025]
    # A second compact table reuses the columns; labels reflect both sources explicitly.
    b=table(s,15,'2025 · sectores IPSS (taxonomía institucional)',['Sector IPSS','Ret. H001','Casos H001','Tasa H001','Ret. SIES','Casos SIES','Tasa SIES','Diferencia (pp)'],[[x['sector'],x['retenidos'],x['casos'],None,*totals[x['sector']],None,None] for x in sectors],'H001/grupos/sector[2025]; SIES_GRANULAR por código + SECTORES; contraste V24_COPIA_RECALCULADA/4_Resultados_Sector!H7:O9',{'B':'count','C':'count','D':'pct','E':'count','F':'count','G':'pct','H':'pp'})
    for i,x in enumerate(sectors):
        r=16+i;sr,sc=totals[x['sector']];formula(b,'D',i,f'=B{r}/C{r}',x['tasa'],s);formula(b,'G',i,f'=E{r}/F{r}',sr/sc,s);formula(b,'H',i,f'=(D{r}-G{r})*100',(x['tasa']-sr/sc)*100,s)
    note(s,21,'C: agregación por sector con mapeo IPSS aprobado. No son áreas SIES/CINE-F. Los 92 «no considerar» SIES están fuera de los 1.163 casos.')
    note(s,24,'E cerrada, 2022/2024: las reglas y la cobertura ya se investigaron; no hay granular oficial comparable que permita acreditar la causa. No se reabre la conciliación.')
    note(s,27,'En 2025 los sectores suman los totales de cada fuente. Las tasas institucionales se obtienen de conteos, nunca del promedio de tasas sectoriales.')

    s=sheet('2_Intertrimestral_2026','Seguimiento T1 a T2/T3 · 2026',[30,21,17,17,16,18,13,56])
    note(s,3,'BORRADOR · indicador interno · no oficial · paridad con fuentes crudas pendiente')
    v=cache['18_Foto_Intertrimestral']; rows=[]
    specs=[('T1 → T2','Proxy actividad',v['E11'].value,v['D11'].value,'Personas','C/E','Estado del extracto PROMEDIOS; no acredita foto de matrícula.'),('T1 → T3','Foto T3',v['G93'].value,v['F93'].value,'Personas','B/C/D','1.471 originales menos 4 NO APLICA. T3 en curso al 22-10-2026.'),('Ingreso T2 → T2','Renuncias agregadas',v['G94'].value,v['F94'].value,'Matrículas','B/C/E','528 matrículas y 527 personas; no dividir por personas.'),('Ingreso T2 → T3','Foto T3',v['G12'].value,v['D12'].value,'Personas','B/C/E','Denominador de personas de la foto; no confundir con la tasa T2.'),('Ingreso T3 → T3','Foto T3',v['G13'].value,v['D13'].value,'Personas','B/C/E','Ingreso T3; T2 no aplica. Foto en curso.'),('Total único → T3','Foto T3',v['G95'].value,v['F95'].value,'Personas','C/D/E','2.523 originales menos 4 NO APLICA; deduplicación anual.'),('Total único → T2','Sin foto completa',None,None,'Personas','E','No calcular tasa; no equivale a cero.'),('Ingreso T3 → T2','No aplica',None,None,'Personas','C/E','No pertenece al seguimiento previo a su ingreso.')]
    for label,kind,n,d,u,c,l in specs: rows.append([label,kind,n,d,None,u,c,l])
    b=table(s,6,'Medidas documentadas en la regla vigente (borrador)',['Medida','Evidencia','Numerador','Denominador','Tasa','Unidad','Clase','Límite'],rows,'INTER/reglas y tipos_de_celda; BASES_V16; V24_COPIA_RECALCULADA/18_Foto_Intertrimestral!D11:H14,D93:I95',{'C':'count','D':'count','E':'pct'})
    for i,(_,_,n,d,*_) in enumerate(specs):
        if n is not None: formula(b,'E',i,f'=C{7+i}/D{7+i}',n/d,s)
    # Compare all numeric rates with the already-governed recalculation.
    for i,cell in enumerate(['F11','H93','H94','H12','H13','H95']): assert abs(specs[i][2]/specs[i][3]-v[cell].value)<1e-12
    note(s,18,'No sumar estas filas: grupos y totales se solapan. Las matrículas T2 no se agregan a las personas. «Sin foto completa» y «No aplica» no son ceros.')
    note(s,21,'E: cinco personas con dos matrículas vigentes están en la construcción desde fuentes y no en V24; la causa de exclusión sigue pendiente. Se conserva V24.')
    note(s,24,'E: faltan la paridad completa de proxies y desgloses sector/tipo y la auditoría integral de exclusión de diplomados en V24. No se activa el indicador.')

    par=js('H002'); rev=js('REVISION'); s=sheet('3_Paridad_SIES_2027','Paridad de códigos · catálogo 2027',[27,36,44,44,20,16,14,20]);s['freeze']=6
    note(s,3,'Catálogo agregado, no asignación de personas ni archivo de carga. Oferta V06 preparada (C); regla interna activa (D).')
    states=sorted(set(x['estado'] for x in par['filas']))
    b=table(s,172,'Resumen por estado · unidad: códigos por año',['Estado','Códigos 2026','Códigos 2027'],[[st,sum(x['anio']==2026 and x['estado']==st for x in par['filas']),sum(x['anio']==2027 and x['estado']==st for x in par['filas'])] for st in states],'H002/filas agrupadas por anio y estado',{'B':'count','C':'count'})
    b['rows'].append(['TOTAL',None,None])
    for col,year in [('B',2026),('C',2027)]:
        formula(b,col,len(states),f'=SUM({col}173:{col}{172+len(states)})',sum(x['anio']==year for x in par['filas']),s)
    # Only summary count remains sourced (both years); full detail is 2027 as requested.
    note(s,183,'ANTIGUO: conserva el código MU 2026. Revisión solo si falta código, no existe en oferta 2027 o tiene VIGENCIA 3; tipo E queda informativo.')
    note(s,186,'NUEVO: única versión VIGENCIA 1 de su familia y tipo de plan UMAS; sin admisión queda informativo. No se elige versión por nombre.')
    note(s,189,'REVISIÓN: tipo desconocido con admisión, incompatibilidad o varias versiones vigentes. Los 37 códigos nuevos se cierran con la MU 2027; no sumar a 34 opciones.')
    detail=[x for x in par['filas'] if x['anio']==2027]
    table(s,6,'Tabla completa 2027 · resumen de estados en fila 172',['Código SIES','Carrera SIES','Plan UMAS / candidatos','Estado','Versión para nuevos','Vigencia 2027','Clase','CODCARPR'],[[x['codigo_unico'],x['nombre_carrera'],'; '.join(x['plan_umas']) or 'E: sin plan',x['estado'],x['version_para_nuevos_2027'],x['vigencia_2027'],x['clase'],'; '.join(x['codcarpr'])] for x in detail],'H002/filas[anio=2027]')
    row=195
    table(s,row,'Revisión real · 34 opciones de catálogo (no personas)',['Familia','Plan UMAS','Motivo','Unidad'],[[x['familia'],x['plan'],x['motivo'],'Opción plan–familia'] for x in rev['nuevos']],'REVISION/nuevos')
    row=235
    pending=[x for x in detail if x['estado']=='NUEVO_2027_SIN_PLAN_E']
    table(s,row,'37 códigos nuevos · cierre cuando exista MU 2027',['Código SIES','Carrera SIES','Forma de cierre','Clase'],[[x['codigo_unico'],x['nombre_carrera'],'MU 2027: cuadratura por código y total','D/E'] for x in pending],'H002/filas; H002_CIERRE/pendientes')
    assert len(detail)==161 and len(pending)==37 and len(rev['nuevos'])==34

    matrix=js('MATRIZ'); gap=js('BRECHA'); s=sheet('4_Brecha_MU_2026','Cuadratura MU · matrículas 2022–2026',[26,15,15,15,15,15,16,17,18])
    note(s,3,'Brecha 2026: CERRADA_DOCUMENTADA, causa E. Decisión D: no continuar con las fuentes disponibles.')
    ms=[(x['anio'],x['mu_sies']) for x in matrix['matriz']]
    b=table(s,6,'Pregrado · publicado menos VIG 1',['Año','Filas base MU','Pregrado MU','Nivel no acreditado','VIG 1 pregrado','VIG 0 pregrado','SIES pregrado','SIES − VIG 1','Control base'],[[y,m['mu_filas'],m['mu_pregrado_filas'],m['mu_nivel_no_acreditado_filas'],m['vig1_pregrado'],m['vig0_pregrado'],m['publicado_pregrado'],None,None] for y,m in ms],'MATRIZ/matriz[*]/mu_sies',{'A':'year',**{x:'count' for x in 'BCDEFGH'},'I':'count'})
    for i,(y,m) in enumerate(ms):
        r=7+i;formula(b,'H',i,f'=G{r}-E{r}',m['diferencia_publicado_menos_vig1'],s);formula(b,'I',i,f'=C{r}+D{r}-B{r}',0,s)
    note(s,14,'2025 cuadra contra VIG 1 + VIG 0 = 2.371: la precarga AC 2026 registra las 73 VIG 0. La causa normativa permanece E; no cambia la regla VIG.')
    b=table(s,18,'Total institucional publicado · nivel separado',['Año','SIES pregrado','SIES postítulo','SIES institucional','Control total'],[[y,m['publicado_pregrado'],m['publicado_postitulo'],m['publicado_institucional'],None] for y,m in ms],'MATRIZ/matriz[*]/mu_sies',{'A':'year','B':'count','C':'count','D':'count','E':'count'})
    for i,(_,m) in enumerate(ms):
        r=19+i;formula(b,'E',i,f'=B{r}+C{r}-D{r}',0,s)
    g=gap['brecha']; b=table(s,27,'2026 · detalle del cierre',['Concepto','Registros','Referencia','Clase'],[['Brecha positiva',g['positiva'],str(g['codigos_positivos'])+' códigos','B/C'],['Brecha negativa',g['negativa'],str(g['codigos_negativos'])+' códigos','B/C'],['Brecha neta',None,'Publicado pregrado − VIG 1 (enlaza H11)','B/C'],['Postítulo registro MRUN',g['postitulo_mrun'],'Agregado; no contiene identificadores','B'],['Postítulo publicado',g['postitulo_publicacion'],'Control frente al registro anterior','B']],'BRECHA/brecha',{'B':'count'})
    formula(b,'B',2,'=H11',g['neta'],s)
    note(s,35,'Reapertura (D): solo con fuente individual SIES de MU 2026, por ejemplo precarga AC 2027 cuando exista; verificar origen, proceso, año y hash.')
    note(s,38,'Cierre documental no significa causa resuelta. La compatibilidad observada con VIG 0 no identifica qué matrículas registró SIES.')

    s=sheet('5_Fuentes','Fuentes fijadas y trazabilidad',[24,94,66,18,20,44]);s['freeze']=6
    note(s,3,'Fecha = verificación de la huella (29-10-2026), no corte de matrícula. Cada bloque indica su referencia temporal. Las rutas son solo metadatos.')
    table(s,6,'Entradas verificadas antes de escribir el libro',['ID','Ruta','SHA-256','Fecha verificación','Clase','Commit / origen'],[[x['id'],x['ruta'],x['sha256'],x['fecha'],x['clase'],x['commit']] for x in src.values()],'FUENTES_FIJADAS.json')
    note(s,28,'H001 activo: commit 4751907; cierre de brecha: 7807542; cierre H002l: 75ed950. Se conservan los hashes completos en la tabla.')
    presentar_h003b(sheets,controls)
    compactar_h003c(sheets,controls)
    model=dict(sheets=sheets,sources=list(src.values()),palette=js('MARCA'),font='Arial',book=BOOK)
    dump(OUT/'MODELO_AGREGADO.json',model);dump(OUT/'CONTROL_CELDAS.json',controls)
    link=OUT/'node_modules'
    if not link.exists(): link.symlink_to(RUNTIME/'node/node_modules',target_is_directory=True)
    autor_h003b()
    margen_texto_xlsx()
    if '--qa' in sys.argv: verificar_h003b()
    cache.close();oficial.close()
    print(json.dumps({'hojas':len(sheets),'controles':len(controls),'fuentes':len(src)},ensure_ascii=False))
if __name__=='__main__': main()
