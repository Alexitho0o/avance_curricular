"""Recalcula copia con LibreOffice es-CL, coteja cada celda, renderiza y entrega.

openpyxl es de lectura solamente. Ajustes OOXML de impresión y cachés conservan
tablas y estilos del original producido por Artifact Tool.
"""
from pathlib import Path
import csv, hashlib, json, math, re, shutil, subprocess, tempfile, zipfile
import xml.etree.ElementTree as ET
import openpyxl, pypdfium2
from pypdf import PdfReader
from construir_excel import OUT, ROOT, RUNTIME, BOOK, verify_sources, sha, dump
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',NS)
def tag(x): return '{'+NS+'}'+x
def replace_zip(p, replacements):
    with zipfile.ZipFile(p) as z: content={n:z.read(n) for n in z.namelist()}
    content.update(replacements)
    q=p.with_suffix('.tmp')
    with zipfile.ZipFile(q,'w',zipfile.ZIP_DEFLATED) as z:
        for n,v in content.items():z.writestr(n,v)
    q.replace(p)
def print_setup(p, layouts):
    replacements={}
    with zipfile.ZipFile(p) as z:
        wb=ET.fromstring(z.read('xl/workbook.xml'))
        calc=wb.find(tag('calcPr'))
        if calc is None: calc=ET.SubElement(wb,tag('calcPr'))
        calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1');calc.set('calcMode','auto')
        names=wb.find(tag('definedNames'))
        if names is None:
            names=ET.Element(tag('definedNames'));wb.insert(list(wb).index(calc),names)
        for old in list(names):
            if old.get('name')=='_xlnm.Print_Area': names.remove(old)
        for i,l in enumerate(layouts):
            dn=ET.SubElement(names,tag('definedName'),{'name':'_xlnm.Print_Area','localSheetId':str(i)})
            dn.text=f"'{l['name']}'!$A$1:${l['lastCol']}${l['lastRow']}"
            name=f'xl/worksheets/sheet{i+1}.xml';root=ET.fromstring(z.read(name))
            pr=root.find(tag('sheetPr'))
            if pr is None: pr=ET.Element(tag('sheetPr'));root.insert(0,pr)
            pp=pr.find(tag('pageSetUpPr'))
            if pp is None: pp=ET.SubElement(pr,tag('pageSetUpPr'))
            pp.set('fitToPage','1')
            for old in ['pageMargins','pageSetup','printOptions']:
                el=root.find(tag(old))
                if el is not None:root.remove(el)
            # These OOXML elements precede tableParts/drawings in the schema.
            idx=next((j for j,e in enumerate(root) if e.tag in [tag('drawing'),tag('legacyDrawing'),tag('tableParts'),tag('extLst')]),len(root))
            elements=[ET.Element(tag('printOptions'),{'headings':'0','gridLines':'0'}),ET.Element(tag('pageMargins'),{'left':'0.25','right':'0.25','top':'0.35','bottom':'0.35','header':'0.1','footer':'0.1'}),ET.Element(tag('pageSetup'),{'paperSize':'8','orientation':'landscape','fitToWidth':'1','fitToHeight':'0'})]
            for el in elements:root.insert(idx,el);idx+=1
            replacements[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        replacements['xl/workbook.xml']=ET.tostring(wb,encoding='utf-8',xml_declaration=True)
    replace_zip(p,replacements)
def run_lo(p, work, fmt):
    profile=work/'perfil';(profile/'user').mkdir(parents=True,exist_ok=True)
    (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale" oor:op="fuse"><value>es-CL</value></prop><prop oor:name="ooLocale" oor:op="fuse"><value>es-CL</value></prop></item></oor:items>''')
    target=work/('recalc' if fmt=='xlsx' else 'pdf');target.mkdir(exist_ok=True)
    cmd=[str(RUNTIME/'bin/override/soffice'),'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to',fmt,'--outdir',str(target),str(p)]
    result=subprocess.run(cmd,text=True,capture_output=True,timeout=120)
    if result.returncode:raise RuntimeError(result.stdout+result.stderr)
    output=target/(p.stem+'.'+fmt)
    if not output.exists():raise RuntimeError(result.stdout+result.stderr)
    return output, result.stdout+result.stderr
def inject_caches(original, recalculated):
    repl={}
    with zipfile.ZipFile(original) as a,zipfile.ZipFile(recalculated) as b:
        for name in a.namelist():
            if not re.fullmatch(r'xl/worksheets/sheet\d+.xml',name):continue
            root=ET.fromstring(a.read(name)); other=ET.fromstring(b.read(name)); by={x.attrib['r']:x for x in other.iter(tag('c'))}
            for c in root.iter(tag('c')):
                if c.find(tag('f')) is None: continue
                calc=by[c.attrib['r']]; v=calc.find(tag('v'))
                if v is None:raise ValueError('Fórmula sin caché '+name+' '+c.attrib['r'])
                old=c.find(tag('v'))
                if old is not None:c.remove(old)
                c.append(ET.fromstring(ET.tostring(v)))
                if calc.get('t')=='e':raise ValueError('Error de fórmula')
                if calc.get('t') in ('str','b'): c.set('t',calc.get('t'))
                elif 't' in c.attrib:del c.attrib['t']
            repl[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    replace_zip(original,repl)
def main():
    verify_sources(); book=OUT/BOOK;layouts=json.loads((OUT/'DISENO.json').read_text())
    print_setup(book,layouts)
    work=Path(tempfile.mkdtemp(prefix='h003_qa_'))
    recalc,log=run_lo(book,work,'xlsx')
    (OUT/'RECALCULO_LIBREOFFICE.log').write_text(log)
    w=openpyxl.load_workbook(recalc,read_only=False,data_only=True)
    f=openpyxl.load_workbook(book,read_only=False,data_only=False)
    controls=json.loads((OUT/'CONTROL_CELDAS.json').read_text()); comparisons=[]
    for x in controls:
        got=w[x['hoja']][x['celda']].value; expected=x['esperado']
        # A literal empty string is serialized by Calc as an empty cell, not zero.
        same=math.isclose(got,expected,rel_tol=1e-12,abs_tol=1e-10) if isinstance(expected,(float,int)) and isinstance(got,(float,int)) else (got==expected or (expected=='' and got is None))
        comparisons.append({**x,'obtenido':got,'coincide':same})
    failures=[x for x in comparisons if not x['coincide']]
    dump(OUT/'DIFERENCIAS_QA.json',failures)
    if failures:dump(OUT/'DIFERENCIAS_QA.json',failures);raise ValueError(f'{len(failures)} diferencias de celda')
    errors=[f'{s.title}!{c.coordinate}' for s in w for row in s for c in row if c.data_type=='e']
    assert not errors, errors
    formulas=sum(c.data_type=='f' for s in f for row in s for c in row)
    assert all(s.sheet_state=='visible' for s in f)
    assert len(f.sheetnames)==6
    tables=sum(len(s.tables) for s in f)
    assert tables==sum(len(s['blocks']) for s in json.loads((OUT/'MODELO_AGREGADO.json').read_text())['sheets'])
    assert all(s.freeze_panes for s in f)
    assert w['2_Intertrimestral_2026']['C13'].value is None and w['2_Intertrimestral_2026']['E13'].value is None
    # Source data are allowlisted aggregate fields. Scan cell text for identifier values.
    patterns={'documento_con_dv':r'(?<![A-Za-z0-9])\d{1,2}[.]\d{3}[.]\d{3}-[0-9Kk](?!\w)|(?<![A-Za-z0-9])\d{7,8}-[0-9Kk](?!\w)', 'correo':r'\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b','telefono':r'\+56\s?[29]\s?\d{4}\s?\d{4}'}
    privacy=[]
    for s in f:
        for row in s:
            for c in row:
                if isinstance(c.value,str):
                    for kind,p in patterns.items():
                        if re.search(p,c.value): privacy.append({'hoja':s.title,'celda':c.coordinate,'tipo':kind})
    assert not privacy, privacy
    # No fields naming individual records are exported; route/hash metadata are allowed.
    assert not any(str(c.value).upper() in ['RUT','CODCLI','MRUN','N_DOC','NOMBRE_ALUMNO','NOMBRES','FECHA_NACIMIENTO'] for s in f for row in s for c in row)
    inject_caches(book,recalc)
    pdf,logpdf=run_lo(book,work,'pdf');shutil.copy2(pdf,OUT/'VISTA_IMPRESION.pdf')
    doc=pypdfium2.PdfDocument(pdf);reader=PdfReader(pdf);render=OUT/'renders';render.mkdir(exist_ok=True)
    pages=[]
    for i,page in enumerate(doc):
        p=render/f'pagina_{i+1:02}.png';page.render(scale=1.2).to_pil().save(p)
        pages.append({'pagina':i+1,'png':str(p),'texto_inicial':reader.pages[i].extract_text()[:100]})
    text='\n'.join(p.extract_text() for p in reader.pages)
    assert '73,4%' in text and '92,7%' in text and '1.467' in text and '2.519' in text, 'Formato es-CL no acreditado en PDF'
    assert '####' not in text
    dump(OUT/'PAGINAS_RENDER.json',pages)
    # Confirm preservation of the starting working tree and every pinned source.
    initial=json.loads((OUT/'ESTADO_INICIAL.json').read_text())
    assert subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()==initial['head']
    assert all(sha(ROOT/p)==h for p,h in initial['archivos_ajenos'].items())
    verify_sources()
    target=Path.home()/'Desktop/ENTREGA_RETENCION_20261031';target.mkdir(exist_ok=True)
    shutil.copy2(book,target/BOOK);assert sha(book)==sha(target/BOOK)
    qa={'estado':'PASS_TECNICO_PENDIENTE_REVISION_VISUAL_Y_APERTURA_EXCEL','fuentes_hash':len(verify_sources()),'controles_celda':len(comparisons),'diferencias_celda':0,'formulas':formulas,'errores_formula':errors,'hojas':len(f.sheetnames),'tablas_autofiltro':tables,'hojas_ocultas':0,'privacidad':privacy,'paginas_png':len(pages),'formato_chileno_pdf':'PASS','ajenos_preservados_archivos':len(initial['archivos_ajenos']),'sha256_libro':sha(book),'copia_escritorio':str(target/BOOK),'sin_commit_push':True,'motor':'LibreOfficeDev 26.8 headless, perfil es-CL aislado','tipografia':'Arial; sustitución declarada'}
    dump(OUT/'QA.json',qa);dump(OUT/'CONTROL_RESULTADOS.json',comparisons)
    with (OUT/'CONTROL_RESULTADOS.csv').open('w',newline='') as out:
        wr=csv.DictWriter(out,fieldnames=['hoja','celda','esperado','fuente','tipo','formula','obtenido','coincide']);wr.writeheader();wr.writerows(comparisons)
    md=['# H003 · Informe de entrega','',f'Libro: `{BOOK}`. SHA-256: `{sha(book)}`.','',f'{len(comparisons)} celdas contrastadas; cero diferencias; {formulas} fórmulas recalculadas sin errores. {tables} tablas con filtro.','', 'No se recalculan H001 ni reconciliaciones cerradas. Se extraen sus agregados. Las tasas conservan precisión subyacente; la visualización usa un decimal.','', 'Copia en Escritorio idéntica. Fuentes y 24 entradas ajenas preservadas (25 archivos). Sin commit/push.','', '## Límites visibles','', '- H001: causas 2022/2024 cerradas E, no resueltas.','- Intertrimestral: borrador; paridad cruda, proxies y desgloses pendientes; no mezclar personas/matrículas.','- H002: 34 opciones y 37 códigos, solapados; oferta V06 preparada C.','- MU 2026: cierre documental de 226; reabrir con fuente individual SIES.','- Marca: paleta oficial; Arial sustituye fuentes oficiales no instaladas.','', '## Control de cada cifra numérica del libro','', '| Hoja y celda | Valor | Fuente / origen | Coincide |','|---|---:|---|---|']
    for x in comparisons:
        if isinstance(x['esperado'],(int,float)):
            md.append(f'| {x["hoja"]}!{x["celda"]} | {x["obtenido"]} | {x["fuente"]} | Sí |')
    md+=['','El control completo, incluidos textos del catálogo y fuentes, está en CONTROL_RESULTADOS.csv/.json.','', '## Renderizado','']+[f'- {x["png"]}' for x in pages]
    (OUT/'INFORME.md').write_text('\n'.join(md)+'\n')
    dump(OUT/'MANIFEST.json',{'head':initial['head'],'fuentes':'FUENTES_FIJADAS.json','salidas':{p.name:sha(p) for p in [book,OUT/'CONTROL_RESULTADOS.json',OUT/'QA.json',OUT/'VISTA_IMPRESION.pdf']}})
    print(json.dumps(qa,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
