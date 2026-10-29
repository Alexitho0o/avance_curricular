#!/usr/bin/env python3
"""H002g: cohortes observadas y propuestas agregadas; ninguna adjudicación personal."""
import argparse,csv,io,json,re,sys,itertools,hashlib
from pathlib import Path
from collections import Counter,defaultdict
sys.dont_write_bytecode=True
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent;DATA=ROOT/'data/paridad_codigo_sies/cohorte_2007_2026'
sys.path.insert(0,str(ROOT/'outputs/paridad_anual_20261029'))
from construir_anual import sha,js,read,txt,number,ident,umas_ident,family,career,selected,sheet_rows,codigo
PRIVATE=Path.home()/'.local/share/ipss/h002g_20261029'

def cohort(y,s,observed):
    y=number(y);return (y,txt(s)) if y and 1900<y<=observed and txt(s) in ('1','2') else (None,'SIN_DATO')

def origin_observed(year,semester):
    # La cohorte de origen se informa literal, incluidos valores especiales;
    # no interviene en la llave de estabilidad de carrera actual.
    return number(year),txt(semester)

def transition(a,b):
    return 'MANTIENE' if len(a)==len(b)==1 and a==b else ('CAMBIA' if len(a)==len(b)==1 else 'MULTIPLE_E')

def choose_da(candidates,cy,sem,year):
    candidates=set(candidates)
    if len(candidates)>1 and cy:
        byyear={x for x in candidates if x[4]==cy}
        if len(byyear)==1:candidates=byyear
        elif year==2026:
            bysem={x for x in byyear if {'1':'1','2':'2','3':'2'}.get(x[5])==sem}
            if len(bysem)==1:candidates=bysem
    return next(iter(candidates)) if len(candidates)==1 else None

def resolve_da(candidates,cy,sem,year,active):
    chosen=choose_da(candidates,cy,sem,year)
    if chosen:return chosen,'IDENTIDAD_FAMILIA_CONTEXTO'
    # Actividad sólo desempata; su ausencia no invalida un expediente único.
    narrowed={c for c in candidates if (c[0],c[1]) in active}
    chosen=choose_da(narrowed,cy,sem,year)
    return chosen,'ACTIVIDAD_ANUAL_DESAMBIGUA' if chosen else 'SIN_UNICO'

def solve(plans,targets,compat):
    """Asignación íntegra de cada grupo a un código: nunca divide grupos para cuadrar."""
    keys=sorted(plans);options=[sorted(compat.get(k,[])) for k in keys]
    if not keys or any(not x for x in options):return 'NO_CUADRA',[],0
    size=1
    for x in options:size*=len(x)
    if size>100000:return 'AMBIGUA',[],size
    exact=[];near=[]
    for choices in itertools.product(*options):
        got=Counter()
        for k,c in zip(keys,choices):got[c]+=plans[k]
        if any(c not in targets for c in got):continue
        delta={c:got[c]-n for c,n in targets.items()}
        item={'asignacion':dict(zip(keys,choices)),'totales_internos':dict(got),'diferencias':delta}
        if all(v==0 for v in delta.values()):exact.append(item)
        elif all(n>0 and abs(delta[c])/n<=.1 for c,n in targets.items()):near.append(item)
    choices=exact if exact else near
    return ('AMBIGUA' if len(choices)>1 else 'CUADRA_EXACTO' if exact else 'APROXIMA' if near else 'NO_CUADRA'),choices,size

def csvbytes(rows):
    if not rows:return ''
    s=io.StringIO();w=csv.DictWriter(s,fieldnames=list(rows[0]),delimiter=';',lineterminator='\n');w.writeheader()
    for r in rows:w.writerow({k:json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(list,dict)) else '' if v is None else v for k,v in r.items()})
    return s.getvalue()

def build():
    entries=read(OUT/'FUENTES.json')
    for x in entries:
        if sha(x['ruta'])!=x['sha256']:raise ValueError('Hash alterado: '+x['id'])
    src={x['id']:Path(x['ruta']) for x in entries}
    groups=read(src['data/versiones_por_cohorte/20261028/grupos_existentes.json']);master=read(src['data/versiones_historicas_sies_umas/fase_5_20261015/maestro.json'])
    profiles=read(src['outputs/paridad_codigo_sies_20261029/PERFIL_PLANES_H002C.json'])['perfiles']
    fam=defaultdict(set)
    for g in groups:
        for f in g['familias']:fam[f].add((g['codcarpr'],g['sede'],g['jornada']))
    for g in master['equivalencias_programa_2026']:fam[family(g['codigo_unico'])].add((g['codcarpr'],g['sede_umas'],g['jornada_umas']))
    inverse=defaultdict(set)
    for f,ctxs in fam.items():
        for ctx in ctxs:inverse[ctx].add(f)
    formal=defaultdict(set)
    for p in profiles:
        for v in p['perfil_formal_H112B']:formal[(p['codcarpr'],p['CODPESTUD'])].add((txt(v['semestres']),txt(v['tipo_plan_codigo'])))
    public={};pubmeta={}
    for r in csv.DictReader(src['Matricula_2007_2026_WEB_10_07_2026.csv'].open(encoding='latin1'),delimiter=';'):
        if r['CÓDIGO DE INSTITUCIÓN']!='162':continue
        y=int(r['AÑO'].removeprefix('MAT_'));c=r['CÓDIGO CARRERA'];assert (y,c) not in public
        public[y,c]=int(r['TOTAL MATRÍCULA']);pubmeta[y,c]={'nivel':r['NIVEL GLOBAL'],'primer_anio':number(r['TOTAL MATRÍCULA PRIMER AÑO']),'tipo_plan':r['TIPO DE PLAN DE LA CARRERA'],'duracion_total':r['DURACIÓN TOTAL DE CARRERA']}
    # Proyección: ni fecha de nacimiento ni nombres se retienen; MRUN nunca sale de memoria.
    fields=['cat_periodo','codigo_unico','mrun','anio_ing_carr_act','sem_ing_carr_act','anio_ing_carr_ori','sem_ing_carr_ori','tipo_plan_carr','dur_total_carr','forma_ingreso','nivel_global']
    micro=[{k:r[k] for k in fields} for r in csv.DictReader(src['sies_individual'].open(encoding='utf-8-sig'))]
    mc=Counter();siescounts=Counter();byid=defaultdict(lambda:defaultdict(set));gaps=Counter();new=defaultdict(Counter);codeprofiles=defaultdict(set);first={};steps=Counter();entry_forms=defaultdict(Counter)
    for r in micro:
        y=int(r['cat_periodo']);c=r['codigo_unico'];f=family(c);cy,sem=cohort(r['anio_ing_carr_act'],r['sem_ing_carr_act'],y);oy,os=origin_observed(r['anio_ing_carr_ori'],r['sem_ing_carr_ori'])
        mc[y,c]+=1;siescounts[y,c,cy,sem,oy,os]+=1;first[f]=min(y,first.get(f,y));codeprofiles[y,c].add((r['tipo_plan_carr'],r['dur_total_carr']))
        if cy is None:gaps[y]+=1
        elif r['mrun']:byid[(r['mrun'],f,cy,sem)][y].add(c)
        if cy==y and r['nivel_global']=='Pregrado':new[y,f][c]+=1;entry_forms[y,c][r['forma_ingreso'] or 'SIN_DATO']+=1
    controls=[{'anio':y,'codigo_sies':c,'registro_individual':mc[y,c],'publicacion':public.get((y,c),0),'diferencia':mc[y,c]-public.get((y,c),0)} for y,c in sorted(set(mc)|{k for k in public if k[0]<=2025})]
    assert all(x['diferencia']==0 for x in controls),'No cuadra fuente individual con publicación'
    for (_,f,cy,sem),ys in byid.items():
        years=sorted(ys)
        for a,b in zip(years,years[1:]):steps[f,cy,sem,a,b,'|'.join(sorted(ys[a])),'|'.join(sorted(ys[b])),transition(ys[a],ys[b])]+=1
    stability=[dict(zip(['familia','cohorte_anio','cohorte_semestre','anio_desde','anio_hasta','codigos_desde','codigos_hasta','estado'],k),transiciones=n) for k,n in sorted(steps.items())]
    del byid,micro
    print('SIES: control completo aprobado; identificadores descartados de memoria de estabilidad',flush=True)
    # Extractos históricos proyectados: no se elige una foto por nombre ni se mezcla su estado.
    dafields=['CODCLI','CODCARPR','RUT','SEDE','JORNADA','ANOINGRESO','PERIODOINGRESO','ANOMATRICULA']
    da=defaultdict(set);owners=defaultdict(set);da_lines=defaultdict(list);da_total=0;da_coverage=[];historical_primary=defaultdict(set)
    import openpyxl
    for name,p in sorted(src.items()):
        if not name.startswith('DATOSDEALUMNOS'):continue
        w=openpyxl.load_workbook(p,read_only=True,data_only=True)
        usable='Hoja1' in w.sheetnames and {'CODCLI','CODCARPR','RUT'}.issubset({txt(v) for v in next(w['Hoja1'].iter_rows(values_only=True))})
        w.close()
        if not usable:
            da_coverage.append({'fuente':name,'filas':0,'uso':'SIN_HOJA_PLANA_CON_EXPEDIENTE'})
            continue
        n=0
        for line,r in sheet_rows(p,'Hoja1',dafields):
            n+=1;da_total+=1;pe=umas_ident(r)
            if not r['CODCLI']:continue
            ctx=(r['CODCLI'],r['CODCARPR'],r['SEDE'],r['JORNADA'],number(r['ANOINGRESO']),r['PERIODOINGRESO'])
            da[pe].add(ctx)
            if pe:owners[r['CODCLI']].add(pe)
            da_lines[ctx].append((name,line))
            if name=='DATOSDEALUMNOS_12159.xlsx':historical_primary[r['CODCLI']].add(ctx)
        da_coverage.append({'fuente':name,'filas':n,'uso':'IDENTIDAD_EXPEDIENTE_INGRESO; NO_VIGENCIA'})
    # Cada expediente una vez; duplicados contradictorios se conservan como problema.
    contexts=defaultdict(set)
    for pe,ctxs in da.items():
        for ctx in ctxs:contexts[ctx[0]].add(ctx)
    plans=defaultdict(set);allplans=defaultdict(set);ppfields=['CODCLI','CODCARR','PLAN_DE_ESTUDIO','ANO','RUT','DIG'];pp_identity_rejected=0;active=defaultdict(set)
    for name,p in sorted(src.items()):
        if not name.startswith('PROMEDIOSDEALUMNOS'):continue
        n=0
        for _,r in sheet_rows(p,'Hoja1',ppfields):
            n+=1;y=number(r['ANO']);pl=r['PLAN_DE_ESTUDIO'];ca=r['CODCARR'];pe=umas_ident(r)
            if not pe or owners[r['CODCLI']]!={pe}:
                pp_identity_rejected+=1;continue
            if y:active[pe,y].add((r['CODCLI'],ca))
            if y and ca not in ('','NULL') and pl not in ('','NULL'):
                plans[r['CODCLI'],ca,y].add(pl);allplans[r['CODCLI'],ca].add(pl)
        print('PROMEDIOS leído:',name,n,flush=True)
    def plan(ctx,y):
        ps=plans[ctx[0],ctx[1],y]
        if len(ps)==1:return next(iter(ps)),'ANUAL'
        if len(ps)>1:return 'PLAN_MULTIPLE_E','VARIOS_EN_ANIO'
        return 'PLAN_NO_OBSERVADO','SIN_ACTIVIDAD_ANUAL; ENLACE_EXPEDIENTE_CONSERVADO'
    # Catálogo interno de ingresos: ni año del nombre del plan ni última matrícula sustituyen ANOINGRESO.
    ingress=defaultdict(lambda:defaultdict(set));ing_detail=Counter();incomplete=Counter();firstinternal={}
    for cli,ctxs in historical_primary.items():
        if len(ctxs)!=1:incomplete['EXPEDIENTE_CONTEXTO_CONFLICTIVO']+=1;continue
        ctx=next(iter(ctxs));_,ca,s,j,cy,period=ctx
        if not cy or not 1900<cy<=2026:incomplete['SIN_ANIO_INGRESO']+=1;continue
        fs=inverse[ca,s,j]
        for f in fs:firstinternal[f]=min(cy,firstinternal.get(f,cy))
        if len(fs)!=1:incomplete['FAMILIA_AUSENTE_O_MULTIPLE']+=1;continue
        f=next(iter(fs));pl,why=plan(ctx,cy);key='|'.join([ca,pl,s,j]);ingress[cy,f][key].add(cli);ing_detail[cy,f,key,period]+=1
    hist=[];fixed={}
    for y,f in sorted({k for k in new if k[0]<=2021}|{k for k in ingress if 2007<=k[0]<=2021}):
        target=dict(new[y,f]);ints={k:len(v) for k,v in ingress[y,f].items()};compat={};profileflags={}
        for k in ints:
            ca,pl,_,_=k.split('|');pf=formal[ca,pl];options=[]
            for c in target:
                cp=codeprofiles[y,c];expected={(d,'3' if 'Continuidad' in t else '1') for t,d in cp}
                if not pf or pf & expected:options.append(c)
            compat[k]=options;profileflags[k]='PERFIL_FORMAL_COMPATIBLE' if pf else 'PERFIL_NO_ACREDITADO'
        state,solutions,size=solve(ints,target,compat)
        # No false exact closure if selected SIES cohorts absent in this family/year.
        missing=sum(n for (yy,c,cy,ss,oy,os),n in siescounts.items() if yy==y and family(c)==f and cy is None)
        conditional=bool(missing or any('NO_ACREDITADO'==v.removeprefix('PERFIL_') for v in profileflags.values()) or any('PLAN_NO_OBSERVADO' in k or 'PLAN_MULTIPLE' in k for k in ints))
        rec={'anio':y,'familia':f,'estado':state,'sies_ingreso_actual_anio':target,'interno_ingreso_anio':ints,'periodos_umas':dict(sorted((k+';periodo='+p,n) for (yy,ff,k,p),n in ing_detail.items() if yy==y and ff==f)), 'perfiles':profileflags,'candidatos_por_plan':compat,'asignaciones':solutions,'combinaciones_evaluables':size,'cohorte_sies_ausente_filas':missing,'alcance':'ANUAL; semestre SIES no homologado automáticamente a período UMAS','condicionada_E':conditional,'clase':'C/E'}
        hist.append(rec)
        if state=='CUADRA_EXACTO' and not conditional:
            for k,c in solutions[0]['asignacion'].items():fixed[y,f,k]=c
    # Confirmación MU -> DA: no requiere actividad. Cohorte MU y periodo UMAS siempre separados.
    mup=next(p for name,p in src.items() if name.startswith('BASE_RETENCION_MATRICULA_UNIFICADA'))
    fields=['ANIO_INFORMADO','TIPO_DOC','N_DOC','DV','COD_SED','COD_CAR','JOR','VERSION','VIG','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FOR_ING_ACT']
    mu_counts=Counter();muall=Counter();detail=[];cohort_codes=defaultdict(set);history_internal=defaultdict(lambda:defaultdict(set));coverage=defaultdict(Counter);plan_cov=defaultdict(Counter)
    for line,r in sheet_rows(mup,'BASE_RETENCION_MU',fields):
        y=number(r['ANIO_INFORMADO']);c=codigo(r);f=family(c);pe=ident(r['TIPO_DOC'],r['N_DOC'],r['DV']);inc=selected(y,r['VIG']);cy,sem=cohort(r['ANIO_ING_ACT'],r['SEM_ING_ACT'],y);oy,os=origin_observed(r['ANIO_ING_ORI'],r['SEM_ING_ORI'])
        poss={ctx for ctx in da.get(pe,set()) if pe and (ctx[1],ctx[2],ctx[3]) in fam[f] and owners[ctx[0]]=={pe}}
        why='IDENTIDAD_NO_COMPARABLE' if not pe else 'SIN_EXPEDIENTE_COMPATIBLE'
        chosen,selection_method=resolve_da(poss,cy,sem,y,active[pe,y])
        if chosen:
            cli,ca,s,j,iy,ip=chosen;pl,pwhy=plan(chosen,y);why='ENLACE_UNICO_DA'
        else:cli='';ca='SIN_ENLACE';s='';j='';iy=None;ip='';pl='SIN_ENLACE';pwhy='';why='VARIOS_EXPEDIENTES_COMPATIBLES' if poss else why
        key=(y,c,ca,pl,s,j,cy,sem,oy,os,iy,ip,why,pwhy)
        mu_counts[key,inc]+=1;muall[y,c,inc]+=1;coverage[y]['filas_mu']+=1;coverage[y]['enlace_todas']+=bool(chosen);coverage[y]['seleccionadas']+=inc;coverage[y]['enlace_seleccionadas']+=bool(chosen) and inc
        if inc and chosen:
            cohort_codes[ca,pl,s,j,cy,sem,y].add(c);plan_cov[y][pwhy]+=1
            if cy is not None:history_internal[cli,f,cy,sem][y].add(c)
        detail.append({'fila_fuente_mu':line,'anio':y,'codigo_sies':c,'incluida':inc,'identidad':pe,'candidatos':sorted(poss,key=str),'seleccion':chosen,'plan':pl,'plan_respaldo':pwhy,'motivo':why,'metodo_enlace':selection_method,'cohorte_sies':[cy,sem],'cohorte_origen_sies':[oy,os],'ingreso_umas_vs_mu':'COINCIDE' if chosen and chosen[4]==cy else 'DIFIERE_E' if chosen else 'SIN_ENLACE','filas_DA':da_lines[chosen] if chosen else []})
    # Privado sin MRUN ni fecha de nacimiento; no escribir fuentes SIES individuales.
    PRIVATE.mkdir(parents=True,exist_ok=True);PRIVATE.chmod(0o700);p=PRIVATE/'ENLACES_MU_DA.json';p.write_text(js(detail));p.chmod(0o600)
    matrix=[]
    for (y,c,cy,sem,oy,os),n in sorted(siescounts.items(),key=lambda x:str(x[0])):
        matrix.append({'origen':'SIES_INDIVIDUAL','anio':y,'familia':family(c),'codigo_sies':c,'codcarpr':'NO_ENLAZADO_A_MRUN','plan_umas':'NO_ENLAZADO_A_MRUN','sede_umas':'','jornada_umas':'','cohorte_anio':cy,'cohorte_semestre':sem,'origen_anio':oy,'origen_semestre':os,'ingreso_umas_anio':None,'ingreso_umas_periodo':'','matriculas':n,'excluidas_vig':0,'estado':'COHORTE_OBSERVADA' if cy else 'COHORTE_AUSENTE_E','plan_respaldo':'','comparacion_totales':'NO_APLICA','clase':'B/C'})
    comparisons=Counter()
    for key in sorted({k for k,inc in mu_counts},key=str):
        y,c,ca,pl,s,j,cy,sem,oy,os,iy,ip,why,pwhy=key;n=mu_counts[key,True];ex=mu_counts[key,False];state='SIN_ENLACE' if ca=='SIN_ENLACE' else 'COHORTE_REPARTIDA' if len(cohort_codes[ca,pl,s,j,cy,sem,y])>1 else 'COHORTE_UNICA'
        expected=fixed.get((cy,family(c),'|'.join([ca,pl,s,j])))
        comparison=('CONFIRMA' if c==expected else 'DIFIERE') if expected else 'SIN_BASE_TOTAL_FIJADA';comparisons[comparison]+=n
        matrix.append({'origen':'MU_DATOS_ALUMNOS','anio':y,'familia':family(c),'codigo_sies':c,'codcarpr':ca,'plan_umas':pl,'sede_umas':s,'jornada_umas':j,'cohorte_anio':cy,'cohorte_semestre':sem,'origen_anio':oy,'origen_semestre':os,'ingreso_umas_anio':iy,'ingreso_umas_periodo':ip,'matriculas':n,'excluidas_vig':ex,'estado':state,'plan_respaldo':pwhy,'comparacion_totales':comparison,'clase':'C/E' if state!='COHORTE_UNICA' or 'NO_' in pl else 'B/C'})
    internalsteps=Counter()
    for (_,f,cy,sem),ys in history_internal.items():
        yy=sorted(ys)
        for a,b in zip(yy,yy[1:]):internalsteps[f,cy,sem,a,b,'|'.join(sorted(ys[a])),'|'.join(sorted(ys[b])),transition(ys[a],ys[b])]+=1
    itr=[dict(zip(['familia','cohorte_anio','cohorte_semestre','anio_desde','anio_hasta','codigos_desde','codigos_hasta','estado'],k),transiciones=n) for k,n in sorted(internalsteps.items())]
    # Regla observada por cohorte; sin umbral inferido ni extrapolación de huecos.
    rulecounts=defaultdict(Counter);ruleyears=defaultdict(set);ruleyearcounts=defaultdict(Counter)
    for r in matrix:
        if r['origen']=='MU_DATOS_ALUMNOS' and r['matriculas'] and r['codcarpr']!='SIN_ENLACE' and r['cohorte_anio']:
            k=(r['familia'],r['codcarpr'],r['plan_umas'],r['sede_umas'],r['jornada_umas'],r['cohorte_anio'],r['cohorte_semestre']);rulecounts[k][r['codigo_sies']]+=r['matriculas'];ruleyears[k].add(r['anio']);ruleyearcounts[k][r['anio'],r['codigo_sies']]+=r['matriculas']
    previous=defaultdict(set)
    for r in read(src['data/versiones_por_cohorte/20261016/resultados_cohortes.json']):
        a=r['datos_alumnos'];c=r['resolucion'].get('codigo_propuesto')
        if c:previous[a['CODCARPR'],a['SEDE'],a['JORNADA'],number(a['ANOINGRESO'])].add(c)
    rules=[]
    for k,counts in sorted(rulecounts.items()):
        f,ca,pl,s,j,cy,sem=k;old=previous[ca,s,j,cy];codes=set(counts);contrast='SIN_ANTECEDENTE_UNICO' if len(old)!=1 else 'CONFIRMA' if codes==old else 'CONTRADICE' if codes.isdisjoint(old) else 'PARCIAL_CON_REPARTO'
        rules.append({'familia':f,'codcarpr':ca,'plan_umas':pl,'sede_umas':s,'jornada_umas':j,'cohorte_anio':cy,'cohorte_semestre':sem,'codigos_y_matriculas_anuales':dict(sorted(counts.items())),'por_anio_codigo':[{'anio':y,'codigo':c,'matriculas':n} for (y,c),n in sorted(ruleyearcounts[k].items())],'anios_observados':sorted(ruleyears[k]),'estado':'VERSION_UNICA_OBSERVADA' if len(codes)==1 else 'COHORTE_REPARTIDA_E','contraste_20261016':contrast,'codigos_previos':sorted(old),'alcance':'C/E: conteos matrícula-año; sin interpolar ni adjudicar; ADMP/CADMP pendiente Alexi'})
    # Revalidación por cohorte: contradice sólo reparto persistente en la misma cohorte/plan,
    # no la mera existencia de un plan actual en una versión histórica.
    table=read(src['data/paridad_codigo_sies/20261029/TABLA_PARIDAD_2026_2027.json'])['filas'];valid=[]
    for t in table:
        if t['anio']!=2027 or t['estado'] not in ('CONTINUA','VERSION_MULTIPLE_RESUELTA','APROBADA_D_2027'):continue
        c=t['codigo_unico'];years=[y for (y,cc),n in public.items() if cc==c and y<=2025 and n>0];mr=[r for r in matrix if r['origen']=='MU_DATOS_ALUMNOS' and r['codigo_sies']==c and r['anio']<=2025 and r['matriculas']]
        conflicts=[r for r in rules if c in r['codigos_y_matriculas_anuales'] and len({v['codigo'] for v in r['por_anio_codigo'] if v['anio']<=2025})>1 and r['plan_umas'] not in ('PLAN_NO_OBSERVADO','PLAN_MULTIPLE_E')]
        unknown=sum(r['matriculas'] for r in mr if r['estado']=='SIN_ENLACE' or r['plan_umas'] in ('PLAN_NO_OBSERVADO','PLAN_MULTIPLE_E') or not r['cohorte_anio'] or r['codcarpr'] not in t['codcarpr'] or r['plan_umas'] not in t['plan_umas']);linked=sum(r['matriculas'] for r in mr)-unknown
        st='SIN_ANCLA' if not years else 'CONTRADICE' if conflicts else 'CONFIRMA_PARCIAL' if unknown or not linked else 'CONFIRMA'
        valid.append({'codigo_sies':c,'estado':st,'ultimo_anio_matricula_hasta_2025':max(years) if years else None,'matriculas_anuales_con_plan_y_cohorte':linked,'sin_enlace_o_plan_o_cohorte':unknown,'cohortes_con_reparto':conflicts,'planes_declarados_H002':t['plan_umas'],'carreras_declaradas_H002':t['codcarpr'],'observados_fuera_candidatos_H002':[{'codcarpr':r['codcarpr'],'plan':r['plan_umas'],'cohorte':[r['cohorte_anio'],r['cohorte_semestre']],'anio':r['anio'],'matriculas':r['matriculas']} for r in mr if r['codcarpr']!='SIN_ENLACE' and r['codcarpr'] not in t['codcarpr']],'contraste_2026_sies':public.get((2026,c),0),'alcance':'Coherencia por cohorte del código; no acredita por sí sola todos los planes candidatos H002 ni cambia decisiones D'})
    firstrows=[]
    for f in sorted({family(c) for y,c in public}):
        seq=[]
        for c in sorted({c for y,c in public if family(c)==f}):
            ys=sorted(y for yy,cc in public if cc==c for y in [yy] if public[y,c]>0)
            seq.append({'codigo':c,'primer_anio':min(ys) if ys else None,'perfiles':[dict(anio=y,**pubmeta[y,c]) for y in ys]})
        firstrows.append({'familia':f,'primer_anio_sies':min(y for y,c in public if family(c)==f and public[y,c]>0),'primer_ingreso_umas_observado':firstinternal.get(f),'versiones':seq})
    for h in hist:
        for k,n in h['interno_ingreso_anio'].items():
            ca,pl,s,j=k.split('|');solcodes=sorted({x['asignacion'][k] for x in h['asignaciones']})
            opts=solcodes if solcodes else h['candidatos_por_plan'][k]
            matrix.append({'origen':'CUADRATURA_TOTALES_C','anio':h['anio'],'familia':h['familia'],'codigo_sies':opts[0] if len(opts)==1 else 'OPCIONES_E','codcarpr':ca,'plan_umas':pl,'sede_umas':s,'jornada_umas':j,'cohorte_anio':h['anio'],'cohorte_semestre':'SIN_HOMOLOGACION','origen_anio':None,'origen_semestre':'','ingreso_umas_anio':h['anio'],'ingreso_umas_periodo':'VER_CUADRATURA_PERIODOS','matriculas':0,'excluidas_vig':0,'estado':h['estado'],'plan_respaldo':'ANUAL_SI_DISPONIBLE','comparacion_totales':'CONDICIONADA_E' if h['condicionada_E'] else 'PROPUESTA_C','clase':'C/E','expedientes_internos':n,'opciones_codigo':opts})
    for r in matrix:
        r.setdefault('expedientes_internos',0);r.setdefault('opciones_codigo',[])
    # Balance MU independiente por código; 2026 usa exclusivamente publicación agregada.
    mucontrol=[]
    for y,c in sorted({(y,c) for y,c,inc in muall}|{(y,c) for y,c in public if y>=2022 and pubmeta[y,c]['nivel']=='Pregrado'}):mucontrol.append({'anio':y,'codigo_sies':c,'mu_seleccionada':muall[y,c,True],'mu_excluida':muall[y,c,False],'sies':public.get((y,c),0),'diferencia':public.get((y,c),0)-muall[y,c,True]})
    priorf={r['anio']:r for r in read(src['outputs/paridad_anual_20261029/RESULTADOS.json'])['anual']};mat28={r['anio']:r['cobertura_familia'] for r in read(src['outputs/matriz_cuadratura_totales_20261028/MATRIZ_CUADRATURA_TOTALES_2022_2026.json'])['matriz']}
    annual=[]
    for y,v in sorted(coverage.items()):
        subset=[r for r in matrix if r['origen']=='MU_DATOS_ALUMNOS' and r['anio']==y and r['matriculas']];gc=defaultdict(set)
        for r in subset:
            if r['codcarpr']!='SIN_ENLACE':gc[r['codcarpr'],r['plan_umas'],r['sede_umas'],r['jornada_umas'],r['cohorte_anio'],r['cohorte_semestre']].add(r['codigo_sies'])
        annual.append({'anio':y,**dict(v),'porcentaje_enlace_seleccionadas':round(v['enlace_seleccionadas']/v['seleccionadas']*100,3),'porcentaje_enlace_todas':round(v['enlace_todas']/v['filas_mu']*100,3),'H002f_porcentaje_seleccionadas':priorf[y]['porcentaje_enlazado'],'matriz_28_oct_cobertura_carrera_todas':mat28[y].get('rut_dv_y_carrera',0),'COHORTE_UNICA':sum(len(cs)==1 for cs in gc.values()),'COHORTE_REPARTIDA':sum(len(cs)>1 for cs in gc.values()),'planes':dict(plan_cov[y]),'ingreso_umas_difiere_mu':sum(r['incluida'] and r['seleccion'] is not None and r['ingreso_umas_vs_mu']=='DIFIERE_E' for r in detail if r['anio']==y),'diferencias_codigos':sum(r['diferencia']!=0 for r in mucontrol if r['anio']==y)})
    # Lotes: una entrada familia/problema, no por fragmento anual o sin enlace.
    problems=defaultdict(list)
    for r in hist:
        if r['estado'] in ('AMBIGUA','NO_CUADRA'):problems[r['familia'],r['estado']].append(r)
    for r in stability:
        if r['estado']!='MANTIENE':problems[r['familia'],'CAMBIO_DE_CODIGO_SIES'].append(r)
    for r in itr:
        if r['estado']!='MANTIENE':problems[r['familia'],'CAMBIO_DE_CODIGO_INTERNO'].append(r)
    for r in rules:
        if r['estado']=='COHORTE_REPARTIDA_E':problems[r['familia'],'COHORTE_REPARTIDA'].append(r)
        if r['contraste_20261016']=='CONTRADICE':problems[r['familia'],'DIFIERE_TRABAJO_PREVIO'].append(r)
    for r in valid:
        if r['estado']=='CONTRADICE':problems[family(r['codigo_sies']),'CONTRADICE_H002'].append(r)
    families=sorted({f for f,p in problems},key=lambda f:(next((r['primer_anio_sies'] for r in firstrows if r['familia']==f),9999),f))
    # Agrupa familias de la misma carrera en un lote, no dividir carrera por jornadas.
    cars=[]
    for f in families:
        car=re.sub(r'J\d+$','',f)
        if car not in cars:cars.append(car)
    lots=[]
    for f,p in sorted(problems,key=lambda k:(cars.index(re.sub(r'J\d+$','',k[0])),k[0],k[1])):
        lots.append({'lote':cars.index(re.sub(r'J\d+$','',f))//5+1,'familia':f,'problema':p,'evidencia':problems[f,p],'propuesta':'Conservar E; verificar cohorte y rectificación de versión antes de aplicar una regla fija.' if 'CAMBIO' in p or 'CONTRADICE' in p or 'REPARTIDA' in p else 'Conservar alternativas; acreditar cobertura y perfil antes de fijar paridad por totales.','decision_alexi':''})
    summary={'clases':'B fuente; C implementación; D método Alexi; E pendiente. Ninguna nueva norma A.', 'fuente_sies_filas':sum(mc.values()),'control_sies_publicacion_diferencias':sum(r['diferencia']!=0 for r in controls),'cohorte_actual_incompleta_por_anio':dict(sorted(gaps.items())),'unidad_historico':'familia × año de ingreso, anual; semestre no homologado','historico_estados':dict(Counter(r['estado'] for r in hist)),'historico_condicionado':sum(r['condicionada_E'] for r in hist),'asignaciones_totales_fijadas_no_condicionadas':len(fixed),'anual':annual,'estabilidad_sies_transiciones':dict(Counter({st:sum(r['transiciones'] for r in stability if r['estado']==st) for st in ['MANTIENE','CAMBIA','MULTIPLE_E']})),'estabilidad_interna_transiciones':{st:sum(r['transiciones'] for r in itr if r['estado']==st) for st in ['MANTIENE','CAMBIA','MULTIPLE_E']},'comparacion_totales_matriculas':dict(comparisons),'formas_ingreso_SIES_por_codigo_anio':[{'anio':y,'codigo':c,'conteos':dict(sorted(v.items())),'uso':'B; no homologa vía interna sin relación gobernada'} for (y,c),v in sorted(entry_forms.items())],'validacion_95':dict(Counter(r['estado'] for r in valid)),'regla_fija_confirmada':not any(r['estado']!='MANTIENE' for r in stability),'datos_alumnos':{'fuentes':da_coverage,'filas':da_total,'expedientes':len(contexts),'limite':'Unión de contextos históricos exactos de extractos catalogados, sin sumar fotos ni usar estados. Conflictos de ingreso/período se conservan; no acredita vigencia histórica.'},'filas_PROMEDIOS_sin_identidad_concordante':pp_identity_rejected,'metodos_enlace':dict(Counter(r['metodo_enlace'] for r in detail if r['incluida'] and r['seleccion'])),'historia_interna_fuente':'DATOSDEALUMNOS_12159.xlsx: 15185 filas, mayor extracto plano catalogado; solo ingreso declarado, no estado histórico','historia_interna_exclusiones':dict(incomplete),'lotes':(len(cars)+4)//5,'carreras_lotes':len(cars),'decisiones_reales':len(lots),'ADMP_CADMP':[r for r in matrix if r['origen']=='MU_DATOS_ALUMNOS' and r['familia']=='I162S2C83J4' and r['anio'] in (2025,2026)],'hallazgos_CI':[r for r in matrix if r['origen']=='MU_DATOS_ALUMNOS' and r['anio']==2025 and r['codcarpr'] in ('CIINF','CICRE','CICIB')]}
    objects={DATA/'MATRIZ_PARIDAD_COHORTE.json':{'filas':matrix},DATA/'REGLAS_VERSION_POR_COHORTE.json':{'reglas':rules},OUT/'CONTROL_SIES_PUBLICACION.json':{'filas':controls},OUT/'CONTROL_MU_PUBLICACION.json':{'filas':mucontrol},OUT/'CUADRATURA_2007_2021.json':{'cohortes':hist},OUT/'ESTABILIDAD.json':{'SIES':stability,'MU_interno':itr},OUT/'VALIDACION_95_H002_v2.json':{'codigos':valid},OUT/'RESULTADOS.json':summary,OUT/'LOTES_REVISION_v2.json':{'lotes':lots}}
    files={p:js(v) for p,v in objects.items()};files[DATA/'MATRIZ_PARIDAD_COHORTE.csv']=csvbytes(matrix);files[DATA/'REGLAS_VERSION_POR_COHORTE.csv']=csvbytes(rules);files[DATA/'PRIMER_ANIO_POR_FAMILIA.csv']=csvbytes(firstrows)
    readme='''# H002g · paridad por cohorte 2007–2026\n\nMétodo D de Alexi, 29-10-2026; resultados B/C y pendientes E. Ninguna adjudicación a personas o cambio de H002/v20.\n\nFuentes por hash en MANIFEST. El CSV individual SIES cuadra por código/año con la publicación. Postítulos participan solo del control y del perfil SIES; cuadratura interna histórica usa ingresos pregrado. No hay cruce RUT–MRUN. MRUN solo en memoria; nacimiento no se utiliza.\n\nCada origen de MATRIZ se conserva separado: SIES_INDIVIDUAL, MU_DATOS_ALUMNOS y CUADRATURA_TOTALES_C no se suman. En el tercero se cuenta expedientes_internos, no matrículas. Matrículas-año no equivalen a personas. El total SIES 2026 procede exclusivamente de la publicación; no se inventan cohortes SIES 2026.\n\nCoporte actual ausente queda E, sin sustituir por origen. Períodos UMAS literales; 4/5 no se convierten. La cuadratura 2007–2021 es anual: falta homologación temporal acreditada para transformar períodos históricos en semestres. Las alternativas son propuestas C; total igual no prueba identidad. CUADRA_EXACTO puede ser condicionado por perfil, plan o cobertura faltantes y no se fija para arrastre.\n\nEnlace desde 2022: identidad exacta, familia gobernada y expediente único por sede/jornada/ingreso. PROMEDIOS complementa el plan anual y, solo si hay varios expedientes compatibles, puede desambiguar por actividad del año con identidad exacta. Su ausencia no invalida un expediente ya único. PLAN_NO_OBSERVADO no rompe el enlace de expediente. DA no acredita vigencia histórica; unión de contextos explícitos de extractos catalogados, sin sumar fotos. Conflictos de contexto quedan E. El año/semestre de MU se mantiene separado del ingreso UMAS. Una fecha UMAS posterior no borra un enlace documental de expediente: se informa DIFIERE_E, sin certificar que el plan actual regía entonces. Historia de ingreso usa 12159, el mayor extracto plano catalogado, sin sumar fotos.\n\nReglas: solo cohortes observadas, sin interpolación, versión máxima ni umbrales inventados. Repartos dentro de la misma cohorte se conservan. CONFIRMA H002 es coherencia observada por cohorte del código, no validación universal de todos sus planes; CONFIRMA_PARCIAL expresa falta de plan/enlace o cobertura anual. SIN_ANCLA: sin matrícula hasta 2025; 2026 se informa como contraste.\n\nADMP/CADMP sigue pendiente de Alexi. Los 37 códigos nuevos esperan su primera MU para ejecutar el motor, sin pedir listado a UMAS. La hipótesis de versión fija no se activa si las transiciones la contradicen.\n\nReproducir: python3 outputs/paridad_cohorte_2007_2026/construir_cohortes.py --verificar. Detalle MU privado en ~/.local/share/ipss/h002g_20261029/.\n'''
    files[DATA/'README.md']=readme.replace('Coporte','Cohorte')
    report=['# H002g · resultados por cohorte','', '## Resultado metodológico','', 'B/C: registro SIES = publicación en todos los códigos y años 2007–2025. D: versión fija al ingresar se somete a contraste. E: no se confirma universalmente; cambios observados y datos faltantes permanecen visibles.','', '## ADMP/CADMP · pendiente Alexi','', '| Año | Código | Carrera / plan | Cohorte actual | Matrículas |','|---|---|---|---|---:|']
    for r in summary['ADMP_CADMP']:
        if r['matriculas']:report.append(f"| {r['anio']} | {r['codigo_sies']} | {r['codcarpr']} / {r['plan_umas']} | {r['cohorte_anio']}-{r['cohorte_semestre']} | {r['matriculas']} |")
    report+=['','## Cobertura anual','', '| Año | MU filas | Seleccionadas | % enlace seleccionadas | % H002f | Enlace todas / carrera 28-oct | Única / repartida |','|---|---:|---:|---:|---:|---|---|']
    for r in annual:report.append(f"| {r['anio']} | {r['filas_mu']} | {r['seleccionadas']} | {r['porcentaje_enlace_seleccionadas']} | {r['H002f_porcentaje_seleccionadas']} | {r['enlace_todas']} / {r['matriz_28_oct_cobertura_carrera_todas']} | {r['COHORTE_UNICA']} / {r['COHORTE_REPARTIDA']} |")
    report+=['','La matriz 28-oct acredita persona+carrera; aquí se exige expediente único+sede+jornada+ingreso. No son porcentajes intercambiables. H002f exigía actividad; H002g conserva enlaces sin ella.','', '## Historia y estabilidad','',f"2007–2021 (familia × año de ingreso, comparación anual): {summary['historico_estados']}. Asignaciones no condicionadas para arrastre: {len(fixed)}. Comparación con MU: {dict(comparisons)}.",f"Cohorte actual ausente/inválida: {dict(sorted(gaps.items()))}.",f"SIES, transiciones de un mismo MRUN/familia/cohorte: {summary['estabilidad_sies_transiciones']}. No son personas distintas; incluye observaciones no consecutivas.",f"Interno, transiciones mismo expediente/familia/cohorte MU: {summary['estabilidad_interna_transiciones']}.",'','## Validación H002', '',str(summary['validacion_95']),'','## Primer año por familia','', '| Familia | Primer SIES | Primer ingreso UMAS observado | Versiones: primer año |','|---|---:|---:|---|']
    for r in firstrows:report.append(f"| {r['familia']} | {r['primer_anio_sies']} | {r['primer_ingreso_umas_observado']} | "+'; '.join(v['codigo']+': '+str(v['primer_anio']) for v in r['versiones'])+' |')
    report+=['','## Reglas observadas por familia (sin umbrales inventados)','']
    for f in sorted({r['familia'] for r in rules}):
        rr=[r for r in rules if r['familia']==f];report.append(f"- {f}: "+'; '.join(f"{r['codcarpr']}/{r['plan_umas']} cohorte {r['cohorte_anio']}-{r['cohorte_semestre']} → {r['codigos_y_matriculas_anuales']} ({r['contraste_20261016']})" for r in rr))
    report+=['','## Controles MU por código y año','', '| Año | Código | MU seleccionada | SIES | Diferencia |','|---|---|---:|---:|---:|']
    for r in mucontrol:report.append(f"| {r['anio']} | {r['codigo_sies']} | {r['mu_seleccionada']} | {r['sies']} | {r['diferencia']} |")
    report+=['','Anclas conservadas: 2022/2024/2025; 2023 parcial, C46J1V1 7/6 y C46J4V3 172/173; 2026 contraste, brecha 226 intacta.','',f"Revisión: {summary['lotes']} lotes, {len(cars)} carreras, {len(lots)} problemas familia/tipo. SIN_ENLACE no genera decisiones. Detalle en LOTES_REVISION_v2.json; códigos y años completos en controles JSON."]
    files[OUT/'INFORME.md']='\n'.join(report)+'\n'
    lines=['# H002g · lotes depurados','', 'Una decisión por familia/problema. Conteos de años distintos no son personas únicas. Secuencia por primer año informado; hasta cinco carreras por lote. SIN_ENLACE solo resumen.','']
    for r in lots:
        lines += [f"## Lote {r['lote']} · {r['familia']} · {r['problema']}", '',f"Evidencia agregada: {len(r['evidencia'])} desgloses, conservados íntegros en LOTES_REVISION_v2.json."]
        for e in r['evidencia']:
            if 'asignaciones'in e:lines.append(f"- {e['anio']}: SIES {e['sies_ingreso_actual_anio']}; interno {e['interno_ingreso_anio']}; alternativas {e['asignaciones']}.")
            elif 'transiciones'in e:lines.append(f"- Cohorte {e['cohorte_anio']}-{e['cohorte_semestre']}, {e['anio_desde']}→{e['anio_hasta']}: {e['codigos_desde']}→{e['codigos_hasta']}; {e['transiciones']} transiciones.")
            elif 'codigos_y_matriculas_anuales'in e:lines.append(f"- {e['codcarpr']}/{e['plan_umas']}, cohorte {e['cohorte_anio']}-{e['cohorte_semestre']}: {e['codigos_y_matriculas_anuales']}; previo {e['codigos_previos']}.")
            else:lines.append(f"- {e['codigo_sies']}: {len(e['cohortes_con_reparto'])} repartos de cohorte; {e['matriculas_anuales_con_plan_y_cohorte']} matrículas-año enlazadas.")
        lines+=['', 'Propuesta: '+r['propuesta'], '', 'Decisión Alexi: ____','']
    lines+=['## SIN_ENLACE · sin decisión individual','', '| Año | Matrículas seleccionadas sin expediente único |','|---|---:|']
    for r in annual:lines.append(f"| {r['anio']} | {r['seleccionadas']-r['enlace_seleccionadas']} |")
    files[OUT/'LOTES_REVISION_v2.md']='\n'.join(lines)+'\n'
    files={p:'\n'.join(l.rstrip() for l in s.splitlines()).rstrip()+'\n' if p.suffix=='.md' else s for p,s in files.items()}
    manifest={'fuentes':entries,'salidas':{str(p.relative_to(ROOT)):hashlib.sha256(s.encode()).hexdigest() for p,s in files.items()},'generador':{'ruta':str(Path(__file__).relative_to(ROOT)),'sha256':sha(__file__)},'especificacion_sha256':sha(OUT/'ESPECIFICACION.md'),'auxiliares':{str(p.relative_to(ROOT)):sha(p) for p in [OUT/'test_cohortes.py',OUT/'qa_cohortes.py',OUT/'ESTADO_INICIAL.json',OUT/'FUENTES.json']},'clase':'B/C/D/E','autoriza_adjudicacion':False}
    files[DATA/'MANIFEST.json']=js(manifest)
    return files,summary

def main():
    a=argparse.ArgumentParser();a.add_argument('--verificar',action='store_true');args=a.parse_args();files,s=build()
    if args.verificar:
        bad=[str(p.relative_to(ROOT)) for p,t in files.items() if not p.exists() or p.read_bytes()!=t.encode()]
        if bad:raise SystemExit('Reproducción distinta: '+str(bad))
    else:
        for p,t in files.items():p.parent.mkdir(parents=True,exist_ok=True);p.write_text(t)
    print(js({k:s[k] for k in ['control_sies_publicacion_diferencias','historico_estados','anual','estabilidad_sies_transiciones','estabilidad_interna_transiciones','validacion_95','lotes','decisiones_reales']}))
if __name__=='__main__':main()
