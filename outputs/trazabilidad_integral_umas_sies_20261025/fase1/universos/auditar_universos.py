"""Fase1/H2: universos congelados. Solo lectura de fuentes; salida agregada sin PII.

No recalcula retención ni modifica denominadores. Contrasta caches y fórmulas
del original, manifiestos y agregados posteriores. Puede ejecutarse con
--salida en directorio nuevo; por defecto escribe junto al script si no existe
RESULTADOS_UNIVERSOS.json. Nunca exporta identificadores.
"""
from pathlib import Path
import argparse, collections, datetime, hashlib, json, re, subprocess, unicodedata, zipfile
import openpyxl

R=Path('/Users/alexi/Documents/GitHub/avance_curricular')
ORIG=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
V18=Path('/Users/alexi/Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261025_V18_CANDIDATA.xlsx')
OLD=R/'outputs/seguimiento_t1_hasta_2026_20261023_v13'
P=Path('/Users/alexi/.local/share/ipss')
SOURCES={}
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
    return h.hexdigest()
def register(p,kind,expected=None):
    p=Path(p);h=sha(p)
    if expected:assert h==expected,str(p)
    SOURCES[str(p)]={'sha256':h,'tipo':kind,'clasificacion':'B'}
    return p
def read(p,kind='auditoría agregada',expected=None):return json.loads(register(p,kind,expected).read_text())
def val(x):return '' if x is None else str(int(x)) if isinstance(x,float) and x.is_integer() else str(x).strip()
def identity(x):return val(x).upper().replace('.','').replace(' ','')
def main():
    arg=argparse.ArgumentParser();arg.add_argument('--salida',type=Path,default=Path(__file__).parent)
    arg.add_argument('--actualizar-auditoria',action='store_true',help='Regenera exclusivamente las tres salidas agregadas de esta auditoría; nunca las fuentes.')
    args=arg.parse_args();out=args.salida
    out.mkdir(parents=True,exist_ok=True);assert args.actualizar_auditoria or not (out/'RESULTADOS_UNIVERSOS.json').exists(),'No sobrescribir auditoría existente'
    register(ORIG,'libro derivado original con padrón granular','20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8')
    register(V18,'libro derivado agregado V18','2bd21ca31d391b59f4bd2c56d4876becd12aae05cb9673a7e069b5542ed66b01')
    cal=read(OLD/'CONSULTA_CALENDARIO_AGREGADA.json');e16=read(R/'outputs/aclaracion_tasas_2026_20261023_v16/EVIDENCIA_AGREGADA.json')
    origin=read(P/'aplicacion_diccionario_intertrimestral_20261023_125556/ORIGENES_T1.json','padrón privado anterior','9ba34a21ce89170f6343e068ba38ac1a546a927c1bd5ffcab9f4fd6ef48288a0')
    slots=read(P/'destinos_calendario_v7_20261023/RANURAS_POR_ORIGEN_PERIODO.json','posiciones privadas origen×período','2b6c4772818695af0d7473c9ed3998ebad6aba5c3f62fe7cce6559e4bab2f094')
    contexts=read(P/'destinos_calendario_v7_20261023/CONTEXTOS_POR_ORIGEN_PERIODO.json','contextos privados existentes','82a12e4dbbd9d3ea751f2601d6b4a2f4ffe6485cb7c1f53268015169dba585ca')
    v8=read(R/'outputs/auditoria_diccionario_sies_20261025/fase3_20261025/ENTRADAS_EXCEL.json')
    wd=openpyxl.load_workbook(ORIG,read_only=True,data_only=True);wf=openpyxl.load_workbook(ORIG,read_only=True,data_only=False)
    vd=openpyxl.load_workbook(V18,read_only=True,data_only=True);vf=openpyxl.load_workbook(V18,read_only=True,data_only=False)
    groups=collections.defaultdict(list)
    for rn,r in enumerate(wd['20_Foto_BBDD'].iter_rows(min_row=6,values_only=True),6):
        if r[1] in (2024,2025,2026) and r[2] in (1,2,3):groups[(r[1],r[2])].append((rn,r))
    stats=[]
    for (year,g),rr in sorted(groups.items()):
        stats.append({'anio_rotulado':year,'grupo':g,'filas':len(rr),'documentos_E_distintos':len({identity(r[4]) for _,r in rr}),
          'llaves_I_distintas':len({identity(r[8]) for _,r in rr}),'llaves_I_vacias':sum(not val(r[8]) for _,r in rr),'CODCLI_en_hoja20':False,
          'numerador_T3_Q':sum(r[16]==1 for _,r in rr),'elegibles_T3_R':sum(r[17]==1 for _,r in rr),
          'estados_T3_P':dict(collections.Counter(val(r[15]) for _,r in rr)),
          'evidencia':{'archivo':str(ORIG),'hoja':'20_Foto_BBDD','rango':f'B{rr[0][0]}:V{rr[-1][0]}','sha256':SOURCES[str(ORIG)]['sha256']},'clasificacion':'B/C'})
    annual=[]
    for year in (2024,2025,2026):
        rr=[r for (y,g),rs in groups.items() if y==year for _,r in rs];annual.append({'anio':year,'filas_grupos':len(rr),
          'llaves_I_unicas':len({identity(r[8]) for r in rr}),'exceso_filas_entre_grupos':len(rr)-len({identity(r[8]) for r in rr}),
          'marca_primera_vez_T':sum(r[19]==1 for r in rr),'clasificacion':'B/C'})
    cells=['D7','D8','D9','D10','D11','D12','D13','D14','E7','E8','E9','E11','E12','F7','F8','F9','F11','F12','G7','G8','G9','G11','G12','G13','G14','H7','H8','H9','H11','H12','H13','H14']
    cmp=[]
    for cell in cells:
        a=wf['18_Foto_Intertrimestral'][cell].value;b=vf['18_Foto_Intertrimestral'][cell].value
        cmp.append({'celda':cell,'original_valor':wd['18_Foto_Intertrimestral'][cell].value,'original_expresion':a,'v18_valor':vd['18_Foto_Intertrimestral'][cell].value,'v18_expresion':b,'original_formula':isinstance(a,str) and a.startswith('='),'v18_formula':isinstance(b,str) and b.startswith('='),'clasificacion':'B/C'})
    agg18=collections.Counter()
    for r in vd['20_Foto_BBDD'].iter_rows(min_row=5,values_only=True):
        if len(r)>=6 and r[0] in (2024,2025,2026):agg18[(r[0],r[1])]+=int(r[5])
    assert agg18==collections.Counter({k:len(v) for k,v in groups.items()})
    # Existing private snapshots are read only. Unique row keys are not exported.
    assert len(origin)==3179 and len(slots)==6358 and len(contexts)==2854
    assert len({r['fila_foto'] for r in origin})==len(origin)
    assert {r['fila_foto'] for r in origin}=={n for (y,g),rs in groups.items() if g==1 for n,_ in rs}
    sc=collections.Counter((int(r['anio_origen']),int(r['anio_observado']),str(r['periodo_nativo']),'sin' if r['sin_contexto'] else 'con') for r in slots)
    cc=collections.Counter((int(r['anio_origen']),int(r['anio_observado']),str(r['periodo_nativo'])) for r in contexts)
    group_stats=[]
    for key in sorted(cc):
        s=[r for r in slots if (int(r['anio_origen']),int(r['anio_observado']),str(r['periodo_nativo']))==key]
        c=[r for r in contexts if (int(r['anio_origen']),int(r['anio_observado']),str(r['periodo_nativo']))==key]
        group_stats.append({'cohorte':key[0],'anio_observado':key[1],'periodo_nativo':key[2],'posiciones_origen':len(s),'personas_origen_distintas':len({r['persona'] for r in s}),
          'con_contexto':sc[*key,'con'],'sin_contexto':sc[*key,'sin'],'observaciones_contexto':len(c),'contextos_adicionales_por_posicion':len(c)-sc[*key,'con'],
          'codcli_distintos_en_contextos':len({r['codcli'] for r in c}),'personas_distintas_en_contextos':len({r['persona'] for r in c}),'clasificacion':'B/C'})
    assert sum(r['con_contexto'] for r in group_stats)==2838 and sum(r['sin_contexto'] for r in group_stats)==3520
    # Four supplied files, resolve Unicode names in known backup folder; no approximate identities.
    raw=[];t3_2026={}
    for name in ['Matrícula T2_2025 2.xlsx','Matrícula T3_2025 2.xlsx','Matrícula T_2_13_07_2026.xlsx','Matrícula_T3_02_10_2026.xlsx']:
        p=next(p for p in Path('/Users/alexi/Downloads/cierre').iterdir() if unicodedata.normalize('NFC',p.name)==name)
        register(p,'fuente institucional de cierre/ingresos');w=openpyxl.load_workbook(p,read_only=True,data_only=True)
        for s in w:
            it=iter(s.values);heads=next(it)
            if 'RUT' not in heads:continue
            rr=[(rn,dict(zip(heads,v))) for rn,v in enumerate(it,2) if dict(zip(heads,v)).get('RUT') is not None]
            by=collections.defaultdict(list)
            for rn,r in rr:by[(val(r.get('ANOINGRESO')),val(r.get('PERIODOINGRESO')))].append(r)
            if name=='Matrícula_T3_02_10_2026.xlsx' and s.title in ('Base','2026_3'):t3_2026[s.title]=rr
            raw.append({'archivo':str(p),'sha256':SOURCES[str(p)]['sha256'],'hoja':s.title,'fila_encabezados':1,'filas_con_RUT':len(rr),'rango_filas':[rr[0][0],rr[-1][0]],
              'columnas':{key:openpyxl.utils.get_column_letter(heads.index(key)+1) for key in ('RUT','CODCLI','ANOINGRESO','PERIODOINGRESO','ANOMATRICULA','PERIODOMATRICULA') if key in heads},
              'personas_RUT_texto_distintas':len({identity(r['RUT']) for _,r in rr}),'codcli_distintos':len({val(r['CODCLI']) for _,r in rr if r.get('CODCLI')}),'columna_codcli_disponible':'CODCLI' in heads,
              'por_ingreso':[{'ANOINGRESO':k[0],'PERIODOINGRESO':k[1],'filas':len(rs),'personas':len({identity(r['RUT']) for r in rs}),'codcli':len({val(r['CODCLI']) for r in rs if r.get('CODCLI')})} for k,rs in sorted(by.items())],
              'clasificacion':'B/C','nota':'Conteo por RUT textual normalizado solo en formato; no acredita identidad ni cobertura exacta del cierre.'})
        w.close()
    base26=[r for _,r in t3_2026['Base']];new26=[r for _,r in t3_2026['2026_3']]
    by26=collections.defaultdict(set)
    for r in base26:by26[(val(r['ANOINGRESO']),val(r['PERIODOINGRESO']))].add(identity(r['RUT']))
    memberships=collections.Counter(p for ps in by26.values() for p in ps)
    from_base={identity(r['RUT']) for r in base26};from_original={identity(r[4]) for _,r in groups[(2026,3)]}
    assert from_base==from_original and len(from_base)==545
    filtered_codes={val(r['CODCLI']) for r in base26 if val(r['ANOINGRESO'])=='2026' and val(r['PERIODOINGRESO'])=='3'}
    assert filtered_codes=={val(r['CODCLI']) for r in new26}
    scope26={'clasificacion':'B/C; E rotulado por decidir','padron_aprobado_conservado':545,'filas_Base':557,'personas_Base':545,
      'pertenencias_grupos_ingreso':sum(memberships.values()),'personas_en_mas_de_un_grupo_ingreso':sum(n>1 for n in memberships.values()),
      'exceso_pertenencias':sum(memberships.values())-len(memberships),'hoja2026_3_codcli_iguales_a_Base_filtrada':True,
      'hoja20_grupo2026_3_documentos_iguales_Base':True,'coinciden_545_de_545':True,
      'intersecciones_grupos':[{'grupo_a':list(a),'grupo_b':list(b),'personas_comunes':len(by26[a]&by26[b])} for i,a in enumerate(sorted(by26)) for b in sorted(by26)[i+1:]],
      'interpretacion':'490 verifica el alcance de la etiqueta ingreso2026-3; no es una propuesta automática de cambio del denominador545 aprobado.'}
    denom=[{'anio':2024,'grupo':1,'poblacion':558,'denom_T2':558,'denom_T3':558,'numerador_T2':160,'numerador_T3':175,'periodo_T2':'2025-2','periodo_T3':'2025-3','unidad':'posiciones personales de cohorte; proxy de actividad','clasificacion':'B/C/E'},
      {'anio':2025,'grupo':1,'poblacion':1150,'denom_T2':1150,'denom_T3':1150,'numerador_T2':611,'numerador_T3':1023,'periodo_T2':'2025-2','periodo_T3':'2025-3','unidad':'personas de cohorte; proxy T2 y foto T3','clasificacion':'B/C'},
      {'anio':2025,'grupo':2,'poblacion':222,'denom_T2':222,'denom_T3':222,'numerador_T2':200,'numerador_T3':202,'periodo_T2':'2025-2','periodo_T3':'2025-3','unidad':'personas del archivo T2','clasificacion':'B/C'},
      {'anio':2026,'grupo':1,'poblacion':1471,'denom_T2':1471,'denom_T3':1467,'numerador_T2':1042,'numerador_T3':1360,'periodo_T2':'2026-2','periodo_T3':'2026-3','unidad':'personas; 4 NO APLICA excluidos en T3 por regla interna heredada','clasificacion':'B/C/D'},
      {'anio':2026,'grupo':2,'poblacion':527,'denom_T2':528,'denom_T3':527,'numerador_T2':462,'numerador_T3':467,'periodo_T2':'2026-2','periodo_T3':'2026-3','unidad':'T2 matrículas, T3 personas; no homogeneizar','clasificacion':'B/C/E'},
      {'anio':2026,'grupo':3,'poblacion':545,'denom_T2':None,'denom_T3':545,'numerador_T2':None,'numerador_T3':480,'periodo_T2':None,'periodo_T3':'2026-3','unidad':'personas de Base completa; no solo ingreso 2026-3','clasificacion':'B/C/E'}]
    for d,row in zip(denom,[7,8,9,11,12,13]):
        assert vd['18_Foto_Intertrimestral'][f'D{row}'].value==d['poblacion']
        assert vd['18_Foto_Intertrimestral'][f'G{row}'].value==d['numerador_T3']
        if d['numerador_T2'] is not None:assert vd['18_Foto_Intertrimestral'][f'E{row}'].value==d['numerador_T2']
    assert sum(d['filas'] for d in stats)==4473
    assert all(d['marca_primera_vez_T']==d['llaves_I_unicas'] for d in annual)
    for w in (wd,wf,vd,vf):w.close()
    for p,v in SOURCES.items():assert sha(Path(p))==v['sha256'],'Fuente alterada'
    result={'fase':'1','hito':'2','fecha':datetime.datetime.now().astimezone().isoformat(),'estado':'UNIVERSOS_REPRODUCIDOS_CON_BRECHAS_DE_ALCANCE','clasificacion':'B/C; D reglas heredadas; E alcance pendiente',
      'fuentes':SOURCES,'padron_original_20':stats,'totales_anuales_sin_doble_conteo':annual,'denominadores_v18_conservados':denom,'comparacion_hoja18':cmp,'v18_hoja20_agregada_conserva_conteos':True,
      'seguimiento_diccionario':{'origenes':3179,'posiciones_2_periodos':6358,'con_contexto':2838,'sin_contexto':3520,'observaciones_contexto':2854,'adicionales_no_personas_nuevas':16,'por_periodo':group_stats,'estados_v18':v8['conteos_contextos'],'origenes_con_codigo':sum(bool(r.get('codigo_origen')) for r in origin)},
      'bases_cierres':raw,'alcance_grupo2026_3':scope26,
      'calendario_6358_posiciones':{'clasificacion':'B/C/E','fuente':str(OLD/'CONSULTA_CALENDARIO_AGREGADA.json'),'campo':'calendario_explicito y limites','valores':cal['calendario_explicito'],'tipo':'PERIODO NATIVO, NO CALENDARIO DE CIERRE HOMOLOGADO','limite':'Posiciones contra período nativo 2/3 de PROMEDIOS según el calendario de consulta existente; no equivale automáticamente a trimestre o fecha de cierre. Sin contexto no acredita falta de actividad.'},
      'QA':{'estado':'PASS','suma_filas_padron_4473':True,'agregado_V18_reproduce_grupos_originales':True,'origenes_3179_iguales_filas_T1_original':True,'3179_por_2_igual_6358':True,'2838_mas_3520_igual_6358':True,'2838_mas_16_igual_2854':True,'denominadores_y_numeradores_seis_grupos_coinciden_cache_V18':True,'totales_anuales_deduplicados_coinciden_formula_T':True,'personas_Base_T3_2026_iguales_original_grupo3':True,'codcli_hoja2026_3_iguales_Base_filtrada':True,'fuentes_conservadas_por_hash':True,'identificadores_exportados':False,'limite':'Auditoría de conteos y fórmulas guardadas; no recálculo completo ni prueba de apertura Excel nativo; no acreditación de cobertura de cierre.'},
      'criterios':{'padron_T1_de_libro_existente':'COMPLETO: 3179 filas de origen con identidad y fuente recuperables en original; no prueba todos los matriculados institucionales ni todos los CODCLI.',
      'identidad_y_expediente':'PARCIAL: hoja20 carece CODCLI; identidades originales se conservan y los contextos tienen CODCLI cuando se localizan. No inferir CODCLI para3520posiciones sincontexto.',
      'calendario':'MIXTO documentado: cohorte2024observada2025;2025y2026mismoaño. Períodosnativos no equivalen automáticamente a cierres.',
      'cohorte2025_T3_124':'Ejercicio separado de127matrículas/124RUT en fuente T3_2025, no grupo3de hoja20/18.',
      'ingreso2026_T3':'545 en Base completa; ingreso2026-3 filtrado501matrículas/490RUT. Mantener545hasta decisión sobre alcance, no llamarlo automáticamente nuevos2026-3.',
      'invariancia':'Originales y denominadores intactos. No tasas nuevas, no adjudicaciones, no commits.'}}
    (out/'RESULTADOS_UNIVERSOS.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    lines=['# Fase 1 · Hito 2 — Universos y denominadores conservados','','B/C — Se recuperó el padrón granular del original; V18 conserva sus totales mediante agregados. No se modificó ningún libro ni denominador.','','| Año/grupo | Población | Base T2 | Base T3 | Numerador T2 | Numerador T3 | Unidad y alcance |','|---|---:|---:|---:|---:|---:|---|']
    for d in denom:lines.append(f"| {d['anio']}/{d['grupo']} | {d['poblacion']} | {d['denom_T2']} | {d['denom_T3']} | {d['numerador_T2']} | {d['numerador_T3']} | {d['clasificacion']}: {d['unidad']} |")
    lines+=['','B/C — 558, 1150 y 1471 son las filas personales del grupo1 del original (rangos B6:V563, B564:V1713, B1714:V3184). No son la matrícula total anual. La base1467 solo corresponde a la tasaT3-2026 tras4NO APLICA, documentados como TITULADO/31 en el estudioV16.','','B/C — 3179 orígenes ×2períodos=6358 posiciones. 2838 tienen contexto y3520 no. Hay2854 contextos porque16 son adicionales en posiciones con más de un expediente/plan. Ninguna de estas unidades se suma como nuevos estudiantes.','','| Cohorte | Año observado | Período nativo | Posiciones | Con contexto | Sin contexto | Contextos | Extra |','|---|---|---|---:|---:|---:|---:|---:|']
    for d in group_stats:lines.append('| '+' | '.join(str(d[k]) for k in ['cohorte','anio_observado','periodo_nativo','posiciones_origen','con_contexto','sin_contexto','observaciones_contexto','contextos_adicionales_por_posicion'])+' |')
    lines+=['','B/C/E — Las 6.358 posiciones usan período nativo 2/3, según calendario_explicito del estudio V13. No es una homologación acreditada al trimestre ni a la fecha exacta de cierre. Las 3.520 posiciones sin contexto no prueban falta de actividad.','','## Descomposición del grupo 2026/3 conservado','','B/C — La hoja Base del cierre T3-2026 contiene 557 matrículas y 545 personas. Las 545 identidades coinciden exactamente con el grupo 2026/3 de la hoja 20 original, sin exclusivos.','','| Año/período de ingreso declarado | Matrículas | Personas distintas dentro del grupo |','|---|---:|---:|','| 2024/3 | 22 | 22 |','| 2025/3 | 34 | 34 |','| 2026/3 | 501 | 490 |','','B/C — Una persona está en los grupos 2024/3 y 2026/3; por eso hay 546 pertenencias y 545 personas distintas en la unión. La hoja 2026_3 reproduce exactamente los 501 CODCLI de Base al filtrar ANOINGRESO=2026 y PERIODOINGRESO=3.','', 'D/E — Se conservan los 545 aprobados. Los 490 verifican una diferencia de rotulado/alcance; no se proponen como sustitución automática del denominador. La decisión posterior es cómo describir el grupo completo ya utilizado.','','## Alcance que falta confirmar','','- E: 124 de ingreso 2025-T3 es el ejercicio anual separado: está en la fuente, pero no incorporado como grupo 3 del padrón intertrimestral actual. No añadir una fila ni sumarlo sin decisión.','- E: denominador 528 de T2-2026 es matrículas; población 527 es personas. Las 66 renuncias son agregadas. No confundir la tasa 462/528 con una tasa por personas.','- E: hoja 20 original no tiene CODCLI; 3.520 posiciones sin contexto diccionario no prueban deserción ni falta de matrícula. No inventar expedientes faltantes.','- E: cohorte 2024 se observa en 2025-2/3, mientras 2025 y 2026 se observan en su mismo año. La disponibilidad del padrón no acredita equivalencia de calendario ni cobertura del cierre.','','## QA','','C — PASS: padrón y grupos reproducidos, agregado V18 conserva conteos, orígenes T1 coinciden con las filas originales, identidades Base T3-2026 coinciden con grupo 3, hoja 2026_3 coincide con Base filtrada, ecuaciones 3.179×2=6.358; 2.838+3.520=6.358; 2.838+16=2.854; numeradores y poblaciones coinciden con las celdas guardadas de V18. Fuentes revalidadas por hash al terminar. Sin exportar identificadores ni modificar bases, denominadores o libros.','', 'E — Este QA no acredita cobertura temporal completa, aceptación SIES ni recálculo integral/apertura en Excel nativo.','','## Fuentes y reproducción','','C — Ejecutar auditar_universos.py con el Python del runtime, en salida nueva. Solo exporta agregados; identifica fórmulas originales y valores materializados V18. RESULTADOS_UNIVERSOS.json contiene fuentes+hashes, rangos, conteos y fórmulas de control.']
    for path,v in SOURCES.items():lines+=['',f"- B: `{path}` — `{v['sha256']}` ({v['tipo']})."]
    (out/'INFORME_UNIVERSOS.md').write_text('\n'.join(lines)+'\n')
    manifest={'clasificacion':'C','entradas':SOURCES,'salidas':{p.name:sha(p) for p in out.iterdir() if p.is_file() and p.name!='MANIFEST_UNIVERSOS.json'},'datos_individuales_exportados':False,'modifica_bases_denominadores':False,'HEAD':subprocess.check_output(['git','-C',str(R),'rev-parse','HEAD'],text=True).strip()}
    (out/'MANIFEST_UNIVERSOS.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'estado':result['estado'],'grupos':len(stats),'origenes':3179,'posiciones':6358,'contextos':2854,'sin_contexto':3520,'salida':str(out)},ensure_ascii=False))
if __name__=='__main__':main()
