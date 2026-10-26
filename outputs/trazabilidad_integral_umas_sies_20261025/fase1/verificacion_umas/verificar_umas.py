"""Verifica enlaces ya existentes; no asigna códigos ni exporta identificadores."""
from pathlib import Path
import collections, hashlib, json, re, sqlite3
import openpyxl

ROOT=Path('/Users/alexi/Documents/GitHub/avance_curricular')
OUT=Path(__file__).parent
DET=Path('/Users/alexi/.local/share/ipss/trazabilidad_integral_fase1_20261025/llaves/DETALLE_LLAVES.json')
DB=Path('/Users/alexi/.local/share/ipss/fase3_candidata_20261025/DICCIONARIO_SIES_CANDIDATO.sqlite')
DESC=ROOT/'data/diccionario_sies_seguimiento/v8_20261025_ocho_p/descriptor.json'
def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
    return h.hexdigest()
def val(x):return '' if x is None else str(int(x)) if isinstance(x,float) and x.is_integer() else str(x).strip()
def norm(x):return val(x).upper().replace('.','').replace(' ','')
def mu_key(r):return(val(r['anio']),val(r.get('tipo_doc_original')),norm(r.get('documento_original')),norm(r.get('dv_original')),r['codigo'])

def main():
    assert not (OUT/'RESULTADOS.json').exists(),'Salida existente; no sobrescribir'
    d=json.loads(DET.read_text());detail=d['detalle'];assert len(detail)==2379
    hashes={str(DET):sha(DET),str(DB):sha(DB),str(DESC):sha(DESC)}
    descriptor=json.loads(DESC.read_text())
    def resolve(path):
        declared=Path(path)
        matches=[(Path(p),h) for p,h in descriptor['entradas_sha256'].items() if Path(p).name==declared.name and Path(p).exists()]
        if declared.exists():
            expected=next((h for p,h in descriptor['entradas_sha256'].items() if Path(p).name==declared.name),None)
            if expected:assert sha(declared)==expected
            return declared
        assert len(matches)==1,'Resolver copia congelada por manifiesto'
        actual,expected=matches[0];assert sha(actual)==expected
        return actual
    assert hashes[str(DB)]=='0f073554722a4c0ca9ffc11222aab2c83cd478a10f3bf31aefa8affeeae72c6b'
    c=sqlite3.connect(DB.as_uri()+'?mode=ro',uri=True);c.execute('PRAGMA query_only=ON')
    mi=collections.defaultdict(list);oi=collections.defaultdict(list)
    for person,payload in c.execute('SELECT persona,payload FROM mu_anual'):
        r=json.loads(payload);mi[mu_key(r)].append((person,r))
    for person,year,payload in c.execute("SELECT persona,anio,payload FROM observaciones WHERE periodo='1'"):
        oi[(person,year)].append(json.loads(payload))
    stats=collections.Counter();by=collections.defaultdict(collections.Counter);bycode=collections.defaultdict(collections.Counter);examples={};matching=[]
    for x in detail:
        year=val(x['cohorte']);code=x['codigo_reproducido'];stats['filas_origen']+=1;by[year]['filas_origen']+=1;bycode[code]['filas_origen']+=1
        m=mi[(year,val(x['tipo_documento_origen']),norm(x['documento_original']),norm(x['dv_original']),code)]
        if not m:
            stats['sin_enlace_exacto_indice_mu']+=1;by[year]['sin_enlace_exacto_indice_mu']+=1;bycode[code]['sin_enlace_exacto_indice_mu']+=1;continue
        people={p for p,r in m};assert len(people)==1
        stats['identidad_anio_codigo_completo_mu']+=1;by[year]['identidad_anio_codigo_completo_mu']+=1;bycode[code]['identidad_anio_codigo_completo_mu']+=1
        o=oi[(next(iter(people)),year)]
        matches=[r for r in o if (r['codigo_documental_atribuible'] or r['codigo_interno_adoptado'])==code]
        if not o:state='sin_contexto_nativo1'
        elif not matches:state='contexto_nativo1_sin_codigo_igual'
        elif len(matches)!=1:state='varios_contextos_con_codigo_igual'
        else:state='un_contexto_nativo1_con_codigo_igual'
        stats[state]+=1;by[year][state]+=1;bycode[code][state]+=1
        if state!='un_contexto_nativo1_con_codigo_igual':continue
        r=matches[0];matching.append((x,r))
        comp=re.fullmatch(r'I(\d+)S(\d+)C(\d+)J(\d+)V(\d+)',code).groups()
        stats['contextos_sede_jornada_concordantes']+=val(r['sede_dimension'])==comp[1] and val(r['jornada_sies'])==comp[3]
        stats['contextos_documentales']+=bool(r['codigo_documental_atribuible'])
        stats['contextos_internos_sin_documental']+=not bool(r['codigo_documental_atribuible']) and bool(r['codigo_interno_adoptado'])
        approved=[z for z in r['fuentes_promedios'] if z['estado_ramo']=='APROBADO']
        if not approved:continue
        name='Ejemplo A' if comp[2]=='57' and 'Ejemplo A' not in examples else 'Ejemplo B'+comp[4] if year=='2025' and comp[2]=='1' and comp[3]=='1' and comp[4] in ('1','2') and 'Ejemplo B'+comp[4] not in examples else None
        if name:
            examples[name]=(x,r,approved[0])
    c.close()
    assert stats['un_contexto_nativo1_con_codigo_igual']==1926
    assert set(examples)=={'Ejemplo A','Ejemplo B1','Ejemplo B2'}
    def evidence(z,fields):
        p=resolve(z['fuente']);hashes.setdefault(str(p),sha(p))
        return {**{k:z.get(k) for k in ['hoja','fila']+fields},'fuente_declarada':z['fuente'],'fuente':str(p),'sha256':hashes[str(p)]}
    safe=[]
    for name,(x,r,pm) in sorted(examples.items()):
        code=x['codigo_reproducido'];comp=re.fullmatch(r'I(\d+)S(\d+)C(\d+)J(\d+)V(\d+)',code).groups()
        da=r['contextos_DA'][0]
        mu=next(z for z in r['evidencias_matricula'] if z['codigo']==code and str(z['anio'])==val(x['cohorte']))
        # Re-read exact source rows; compare only in memory, export white-listed fields.
        validations=[];actual_fields={}
        for z in (da,pm,mu):
            actual_path=resolve(z['fuente']);w=openpyxl.load_workbook(actual_path,read_only=True,data_only=True);ws=w[z['hoja']]
            heads=list(next(ws.iter_rows(min_row=1,max_row=1,values_only=True)));row=next(ws.iter_rows(min_row=z['fila'],max_row=z['fila'],values_only=True));row=dict(zip(heads,row))
            if z is not mu:assert val(row['CODCLI'])==r['CODCLI']
            if z is pm:
                assert val(row['PLAN_DE_ESTUDIO'])==r['plan'] and val(row['ANO'])==r['ano'] and val(row['PERIODO'])==r['periodo']
                assert norm(row['RUT'])+'-'+norm(row['DIG'])==r['persona']
                actual_fields['PROMEDIOS']={k:row.get(k) for k in ('CODCARR','CARRERA','JORNADA','REGIMEN','PLAN_DE_ESTUDIO','ANO','PERIODO','NIVEL','CODRAMO','ASIGNATURA','ESTADO','DESCRIPCION_ESTADO','CONVALIDADO')}
            elif z is da:
                assert val(row['CODCARPR'])==r['carrera'] and norm(row['RUT'])==r['persona']
                actual_fields['DA']={k:row.get(k) for k in ('CODCARPR','SEDE','JORNADA','ANOINGRESO','PERIODOINGRESO','ANOMATRICULA','PERIODOMATRICULA')}
            else:
                actual='I162S'+val(row['COD_SED'])+'C'+val(row['COD_CAR'])+'J'+val(row['JOR'])+'V'+val(row['VERSION'])
                assert actual==code and norm(row['N_DOC'])==norm(x['documento_original']) and norm(row['DV'])==norm(x['dv_original'])
                actual_fields['MU']={k:row.get(k) for k in ('ANIO_INFORMADO','TIPO_DOC','COD_SED','COD_CAR','JOR','VERSION','MODALIDAD','VIG','ANIO_ING_ORI','SEM_ING_ORI')}
            validations.append({'fuente':str(actual_path),'hoja':z['hoja'],'fila':z['fila'],'identidad_coincide_en_memoria':True,'CODCLI_coincide_en_memoria':True if z is not mu else 'No existe CODCLI en MU; enlace por identidad completa','campos_contexto_coinciden':True})
            w.close()
        safe.append({'ejemplo':name,'clasificacion':'B/C','cohorte':x['cohorte'],'periodo_nativo':r['periodo'],'codigo_umas':r['carrera'],'plan_umas':r['plan'],'sede_umas':da['SEDE'],'jornada_umas':da['JORNADA'],
          'COD_INST':comp[0],'COD_SED':comp[1],'COD_CAR':comp[2],'JOR':comp[3],'VERSION':comp[4],'codigo_sies_completo':code,
          'libro_original':{'fuente':d['fuente'],'sha256':d['sha256'],'hoja9_fila':x['fila_hoja9'],'hoja11_fila':x['primera_fila_hoja11']},
          'fuente_DA':evidence(da,['CODCARPR','SEDE','JORNADA','ANOINGRESO','PERIODOINGRESO','ANOMATRICULA','PERIODOMATRICULA']),
          'fuente_PROMEDIOS':evidence(pm,['ANO','PERIODO','PLAN_DE_ESTUDIO','estado_ramo','convalidado']),
          'fuente_MU':evidence(mu,['anio','codigo']),
          'estado_atribucion':r['estado_documental'],'regla':r['regla'],'aceptacion_SIES':r['aceptacion_SIES'],'plan_efectivo_periodo':r['plan_efectivo_periodo'],'pendiente':r['requisito_pendiente'],
          'campos_observados_releidos':actual_fields,'validacion_filas_fuentes':validations,'no_se_muestran_identificadores':True})
    for p,h in hashes.items():assert sha(p)==h
    result={'clasificacion':'B/C/E','alcance':'Verificación de enlaces existentes para las2379filas elegibles2022-2025; no nuevas adjudicaciones ni tasa de retención.',
      'metodo':'1) tipo documental+documento+DV+año+código completo exactos en índice mu_anual;2) identidad normalizada de esa misma evidencia, año y período nativo1;3) localizar código adoptado igual y contar contextos. No elegir CODCLI por nombre ni trasladar código entre períodos.',
      'conteos':dict(stats),'por_cohorte':{k:dict(v) for k,v in by.items()},'por_codigo_completo':{k:dict(v) for k,v in bycode.items()},'unidades':{'filas':2379,'codigos_completos':len(bycode),'cohorte_codigo':len({(x['cohorte'],x['codigo_reproducido']) for x in detail}),'COD_CAR':len({re.fullmatch(r'I(\d+)S(\d+)C(\d+)J(\d+)V(\d+)',x['codigo_reproducido']).group(3) for x in detail})},'ejemplos':safe,'fuentes_sha256':hashes,
      'limites':['1926 son enlaces existentes a contexto nativo1; no certificación de plan histórico efectivo ni aceptación SIES.','Sin contexto en este índice no prueba ausencia de UMAS: el diccionario tiene alcance de familias y períodos limitado.','Los18 sin índice MU son un límite del índice v8; el enlace de las2379 filas a MU del libro ya fue reproducido.','La versión se lee del ancla individual anual, no se deduce exclusivamente del nombre, plan ni de totales.','Esta auditoría no acredita por sí sola el calendario de cierre T1/T2/T3.'],
      'QA':{'conteos_sumados':stats['identidad_anio_codigo_completo_mu']+stats['sin_enlace_exacto_indice_mu']==2379,'fuentes_sin_cambios_por_hash':True,'ejemplos_releidos_con_CODCLI_contexto_concordante':True,'datos_individuales_exportados':False,'asignaciones_nuevas':0}}
    (OUT/'RESULTADOS.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    lines=['# Verificación UMAS de los enlaces anuales existentes','','B/C — No se modificaron asignaciones, fuentes ni denominadores. Los conteos siguientes describen el índice v8 disponible; no sustituyen la comprobación anual hoja9→MU del libro, que reprodujo2379/2379.','','| Cohorte | Filas originales | Identidad/año/código en índiceMU | Contexto nativo1 único con mismo código | Sin contexto1 | Contexto1 sin código igual | Sin índiceMU |','|---|---:|---:|---:|---:|---:|---:|']
    for year,v in sorted(by.items()):lines.append('| '+' | '.join(str(x) for x in [year,v['filas_origen'],v['identidad_anio_codigo_completo_mu'],v['un_contexto_nativo1_con_codigo_igual'],v['sin_contexto_nativo1'],v['contexto_nativo1_sin_codigo_igual'],v['sin_enlace_exacto_indice_mu']])+' |')
    lines+=['','E — Los1926 contextos concordantes son enlaces existentes de identidad, expediente, plan, año y período nativo1, con código adoptado igual al anual. La fecha efectiva histórica del plan y la aceptación SIES no se vuelven acreditadas por esta comparación. No se presenta este conteo como un cierre nuevo.','','## Ejemplos reales sin identificadores personales']
    for z in safe:
        lines+=['',f"### {z['ejemplo']}",'',f"B/C — Cohorte{z['cohorte']}, período nativo{z['periodo_nativo']}: UMAS `{z['codigo_umas']}`; plan `{z['plan_umas']}`; sede `{z['sede_umas']}`; jornada `{z['jornada_umas']}` → `{z['codigo_sies_completo']}`. Componentes SIES: institución{z['COD_INST']}, sede{z['COD_SED']}, carrera{z['COD_CAR']}, jornada{z['JOR']}, versión{z['VERSION']}.",f"B/C — Libro original: hoja9 fila{z['libro_original']['hoja9_fila']} → hoja11 fila{z['libro_original']['hoja11_fila']}. Código y expediente ya atribuidos por `{z['regla']}`. No es asignación nueva."]
        for fld in ('fuente_DA','fuente_PROMEDIOS','fuente_MU'):
            e=z[fld];lines+=[f"- B/C: `{e['fuente']}`, hoja `{e['hoja']}`, fila{e['fila']}; SHA-256 `{e['sha256']}`."]
        lines+=[f"E — `{z['pendiente']}`."]
    lines+=['','C — Reproducir con verificar_umas.py en una salida nueva. RESULTADOS.json conserva campos verificables y hashes; no contiene RUT, CODCLI personales ni nombres.','', 'E — Los ejemplos B1/B2 muestran coexistencia de versiones en la misma familia de Informática y año. El código de cada persona se conserva desde su evidencia anual; no se generaliza una versión a toda la carrera.']
    (OUT/'INFORME.md').write_text('\n'.join(lines)+'\n')
    (OUT/'MANIFEST.json').write_text(json.dumps({'clasificacion':'C','entradas_sha256':hashes,'salidas_sha256':{p.name:sha(p) for p in OUT.iterdir() if p.is_file() and p.name!='MANIFEST.json'},'fuentes_modificadas':False,'denominadores_modificados':False,'datos_individuales_exportados':False},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'conteos':stats,'ejemplos':[{'ejemplo':z['ejemplo'],'codigo_umas':z['codigo_umas'],'plan':z['plan_umas'],'codigo_sies':z['codigo_sies_completo']} for z in safe]},ensure_ascii=False))
if __name__=='__main__':main()
