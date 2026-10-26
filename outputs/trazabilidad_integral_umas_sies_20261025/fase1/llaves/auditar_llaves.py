"""Fase1/H3. Auditoría de solo lectura del enlace del libro original, sin adjudicar.
Reproduce U:X y AH; compara candidatos sin modificar bases, tasas ni denominadores.
Datos individuales: salida restringida fuera de Git; JSON/MD públicos solo agregados.
"""
from pathlib import Path
from collections import Counter, defaultdict
import hashlib,json,datetime,subprocess,sqlite3,os
import openpyxl
R=Path('/Users/alexi/Documents/GitHub/avance_curricular')
O=Path(__file__).resolve().parent
P=Path('/Users/alexi/.local/share/ipss/trazabilidad_integral_fase1_20261025/llaves')
SRC=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
EXPECTED='20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8'
NS='11_BASE_RETENCION_MU_RAW'; S9='9_Base_Analitica_Detalle'
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
 return h.hexdigest()
def text(x):
 if x is None:return ''
 if isinstance(x,float) and x.is_integer():return str(int(x))
 return str(x)
def key(row,ix):return '|'.join(text(row[i]) for i in ix)
def code(r):return 'I162S%sC%sJ%sV%s'%tuple(text(r[i]) for i in [11,12,14,15])
def write(p,x,private=False):
 payload=json.dumps(x,ensure_ascii=False,indent=2)+'\n'
 if private:
  if p.exists():
   assert p.read_text()==payload,'Detalle previo distinto; no sobrescribir'
   return
  fd=os.open(p,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
  with os.fdopen(fd,'w') as f:f.write(payload)
 else:p.write_text(payload)
assert O.is_relative_to(R)
assert not P.resolve().is_relative_to(R)
P.mkdir(parents=True,mode=0o700,exist_ok=True);os.chmod(P,0o700)
assert subprocess.run(['git','-C',str(P),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
before=sha(SRC);assert before==EXPECTED,'Original distinto'
w=openpyxl.load_workbook(SRC,read_only=True,data_only=True)
f=openpyxl.load_workbook(SRC,read_only=True,data_only=False)
a=[(n,r) for n,r in enumerate(w[S9].iter_rows(min_row=7,max_row=2385,values_only=True),7)]
m=[(n,r) for n,r in enumerate(w[NS].iter_rows(min_row=5,max_row=10204,values_only=True),5)]
afx=list(f[S9].iter_rows(min_row=7,max_row=2385,values_only=True))
mfx=list(f[NS].iter_rows(min_row=5,max_row=10204,values_only=True))
cat=list(w['8_Mapeo_Sector_IPSS'].iter_rows(min_row=6,max_row=22,min_col=2,max_col=4,values_only=True))
assert len(a)==2379 and len(m)==10200
catalog=defaultdict(list)
for n,r in enumerate(cat,6):catalog[text(r[0]).casefold()].append((n,r))
idx=defaultdict(list);rawkeymismatch=[];formmismatch=[]
for (n,r),fr in zip(m,mfx):
 k=key(r,[0,2,3,12]);cached=text(r[33])
 if k!=cached:rawkeymismatch.append(n)
 if fr[33]!=f'=A{n}&"|"&C{n}&"|"&D{n}&"|"&M{n}':formmismatch.append((NS,n,'AH'))
 idx[k.casefold()].append((n,r))
db=sqlite3.connect(':memory:')
db.execute('create table mu(rownum integer,k text,codigo text,sede text,jornada text,version text,tipo text,modalidad text)')
db.executemany('insert into mu values(?,?,?,?,?,?,?,?)',[(n,text(r[33]).casefold(),code(r),text(r[11]),text(r[14]),text(r[15]),text(r[1]),text(r[13])) for n,r in m])
summary=Counter();byyear=defaultdict(Counter);bygroup=defaultdict(Counter);detail=[];cachemismatch=[];independent_disagreements=[]
allkeysused=set();allmatchedbaserows=set()
for (n,r),fr in zip(a,afx):
 year=text(r[1]);counts=byyear[year]
 formulas={20:f'=IFERROR(VLOOKUP($C{n},\'8_Mapeo_Sector_IPSS\'!$B$6:$D$22,3,FALSE()),"")',21:f'=IFERROR(INDEX(\'{NS}\'!$L$5:$L$10204,MATCH($B{n}&"|"&$H{n}&"|"&$I{n}&"|"&$U{n},\'{NS}\'!$AH$5:$AH$10204,0)),"")',22:f'=IFERROR(INDEX(\'{NS}\'!$P$5:$P$10204,MATCH($B{n}&"|"&$H{n}&"|"&$I{n}&"|"&$U{n},\'{NS}\'!$AH$5:$AH$10204,0)),"")',23:f'=IF(OR(V{n}="",P{n}="",W{n}="",U{n}=""),"","I162S"&V{n}&"C"&U{n}&"J"&P{n}&"V"&W{n})'}
 for c,expected in formulas.items():
  if fr[c]!=expected:formmismatch.append((S9,n,openpyxl.utils.get_column_letter(c+1)))
 catrows=catalog[text(r[2]).casefold()];u=catrows[0][1][2] if catrows else ''
 k='|'.join(text(x) for x in (r[1],r[7],r[8],u));allkeysused.add(k.casefold());cand=idx[k.casefold()]
 first=cand[0][1] if cand else None;sed=first[11] if first is not None else '';ver=first[15] if first is not None else ''
 expectedcode='' if any(text(x)=='' for x in [sed,r[15],ver,u]) else f'I162S{text(sed)}C{text(u)}J{text(r[15])}V{text(ver)}'
 for c,ex in [(20,u),(21,sed),(22,ver),(23,expectedcode)]:
  if text(r[c])!=text(ex):cachemismatch.append({'fila_9':n,'columna':openpyxl.utils.get_column_letter(c+1),'esperado':text(ex),'cache':text(r[c])})
 codes={code(x) for _,x in cand};types={text(x[1]).casefold() for _,x in cand};dims={tuple(text(x[i]) for i in [11,12,14,15]) for _,x in cand}
 samejour=[(i,x) for i,x in cand if text(x[14]).casefold()==text(r[15]).casefold()]
 compatible=[(i,x) for i,x in samejour if text(x[1]).casefold()==text(r[6]).casefold() and text(x[13]).casefold()==text(r[14]).casefold()]
 typedcode=[(i,x) for i,x in compatible if code(x)==expectedcode]
 sq=db.execute('select count(*),count(distinct codigo) from mu where k=?',(k.casefold(),)).fetchone()
 sf=db.execute('select codigo from mu where k=? order by rownum limit 1',(k.casefold(),)).fetchone()
 if sq!=(len(cand),len(codes)) or (sf[0] if sf else None)!=(code(first) if first is not None else None):independent_disagreements.append(n)
 flags={'filas_analiticas':True,'coincidencia_unica':len(cand)==1,'llave_repetida':len(cand)>1,'sin_match':not cand,'varios_codigos_completos':len(codes)>1,'repetidos_mismo_codigo':len(cand)>1 and len(codes)==1,'primero_codigo_coincide_con_codigo_excel':first is not None and code(first)==expectedcode,'codigo_excel_no_coincide_primera_fila_mu':first is not None and code(first)!=expectedcode,'codigo_excel_no_en_ninguna_fila_mu_de_llave':bool(cand) and expectedcode not in codes,'codigo_excel_en_alguna_fila_mu':expectedcode in codes,'primera_jornada_difiere':first is not None and text(first[14]).casefold()!=text(r[15]).casefold(),'primera_modalidad_difiere':first is not None and text(first[13]).casefold()!=text(r[14]).casefold(),'primer_tipo_doc_difiere':first is not None and text(first[1]).casefold()!=text(r[6]).casefold(),'sin_candidato_tipo_jornada_modalidad':not compatible,'codigo_excel_respaldado_tipo_jornada_modalidad':bool(typedcode),'varios_codigos_tras_tipo_jornada_modalidad':len({code(x) for _,x in compatible})>1,'catalogo_nombre_sin_match':not catrows,'catalogo_nombre_varios_codcar':len({text(x[1][2]) for x in catrows})>1,'llave_candidatos_varias_sedes':len({text(x[11]) for _,x in cand})>1,'llave_candidatos_varias_jornadas':len({text(x[14]) for _,x in cand})>1,'llave_candidatos_varias_versiones':len({text(x[15]) for _,x in cand})>1,'llave_candidatos_varios_tipos_doc':len(types)>1,'identidad_o_clave_con_comodin':any(x in k for x in ['*','?','~'])}
 for fl,v in flags.items():
  if v:summary[fl]+=1;counts[fl]+=1;bygroup[(year,text(u),text(r[6]))][fl]+=1
 counts['retenidos']+=int(r[5] or 0);summary['retenidos']+=int(r[5] or 0)
 for i,_ in cand:allmatchedbaserows.add(i)
 detail.append({'fila_hoja9':n,'cohorte':year,'tipo_documento_origen':text(r[6]),'documento_original':text(r[7]),'dv_original':text(r[8]),'llave_excel':k,'cod_car_esperado_por_catalogo':text(u),'jornada_hoja9':text(r[15]),'modalidad_hoja9':text(r[14]),'codigo_cache_hoja9':text(r[23]),'codigo_reproducido':expectedcode,'primera_fila_hoja11':cand[0][0] if cand else None,'codigo_primera_fila_mu':code(first) if first is not None else None,'candidatos':[{'fila':i,'tipo_documento':text(x[1]),'sede':text(x[11]),'carrera':text(x[12]),'jornada':text(x[14]),'version':text(x[15]),'modalidad':text(x[13]),'vig':text(x[32]),'codigo_completo':code(x)} for i,x in cand],'controles':flags})
w.close();f.close();db.close()
for label in flags:
 summary.setdefault(label,0)
 for c in byyear.values():c.setdefault(label,0)
 for c in bygroup.values():c.setdefault(label,0)
allrepeated={k:v for k,v in idx.items() if len(v)>1}
usedrepeated={k:idx[k] for k in allkeysused if len(idx[k])>1}
def dupstats(groups):return {'llaves_repetidas':len(groups),'filas_en_llaves_repetidas':sum(len(v) for v in groups.values()),'con_un_solo_codigo':sum(len({code(r) for _,r in v})==1 for v in groups.values()),'con_varios_codigos':sum(len({code(r) for _,r in v})>1 for v in groups.values()),'copias_exactas_columnas_A_AG_excedentes':sum(len(v)-len({tuple(text(x) for x in r[:33]) for _,r in v}) for v in groups.values())}
after=sha(SRC);assert after==before
private={'clasificacion':'B/C/E','fuente':str(SRC),'sha256':before,'detalle':detail,'discordancias_formula':formmismatch,'discordancias_cache':cachemismatch,'discordancias_AH_cache':rawkeymismatch};write(P/'DETALLE_LLAVES.json',private,True)
report={'clasificacion':'B observado/C implementación/E limitación','fecha':datetime.datetime.now().astimezone().isoformat(),'fuente':{'ruta':str(SRC),'sha256_antes':before,'sha256_despues':after,'hoja_analitica':S9,'rango':'B7:X2385','hoja_mu':NS,'rango_mu':'A5:AH10204'},'universo':{'filas_hoja9':len(a),'filas_MU':len(m),'llaves_mu_distintas':len(idx),'llaves_hoja9_distintas':len(allkeysused),'filas_mu_enlazadas_distintas':len(allmatchedbaserows)},'conteos':dict(summary),'por_cohorte':dict(byyear),'por_cohorte_codcar_tipo_doc':[{'cohorte':k[0],'cod_car':k[1],'tipo_doc':k[2],**v} for k,v in sorted(bygroup.items())],'duplicados_mu_completa':dupstats(allrepeated),'duplicados_llaves_usadas':dupstats(usedrepeated),'perfil_llaves_repetidas_mu':[{'anio':text(v[0][1][0]),'cod_car':text(v[0][1][12]),'filas':len(v),'codigos':sorted({code(r) for _,r in v}),'jornadas':sorted({text(r[14]) for _,r in v}),'versiones':sorted({text(r[15]) for _,r in v}),'vigs':sorted({text(r[32]) for _,r in v}),'tipos_doc':sorted({text(r[1]) for _,r in v}),'en_universo_hoja9':k in allkeysused} for k,v in sorted(allrepeated.items())],'QA':{'original_intacto':before==after,'formulas_evaluadas_hoja9':len(a)*4,'formulas_evaluadas_AH':len(m),'formulas_fuera_del_patron':len(formmismatch),'cache_reproduccion_U_V_W_X_distinto':len(cachemismatch),'cache_AH_distinto':len(rawkeymismatch),'contraste_sqlite_independiente_discrepancias':len(independent_disagreements),'conserva_denominador':len(a)==2379,'conserva_retenidos':summary['retenidos']==1476},'detalle_privado':{'ruta':str(P/'DETALLE_LLAVES.json'),'sha256':sha(P/'DETALLE_LLAVES.json'),'modo':'0600','fuera_git_verificado':True},'script':{'ruta':str(Path(__file__).resolve()),'sha256':sha(Path(__file__).resolve())},'limites':['Audita el libro original y su selección por primeras coincidencias; no prueba recepción/aceptación SIES.','No modifica fórmulas ni adopta códigos alternativos.','Coincidencias tipo/jornada/modalidad no acreditan por sí solas atribución CODCLI/plan.','Los conteos de flags se solapan; no sumarlos como grupos excluyentes.','Los indicadores de denominador y retenidos se conservan; no se recalcula elegibilidad ni retención.']}
write(O/'RESULTADOS_AGREGADOS.json',report)
q=report['QA'];good=all([q['original_intacto'],not q['formulas_fuera_del_patron'],not q['cache_reproduccion_U_V_W_X_distinto'],not q['cache_AH_distinto'],not q['contraste_sqlite_independiente_discrepancias'],q['conserva_denominador'],q['conserva_retenidos']])
write(O/'QA.json',{'clasificacion':'C','estado':'CORRECTO' if good else 'OBSERVACIONES_REQUIEREN_REVISION','controles':q,'no_cambios_fuentes_bases_excel_diccionario_denominadores':True})
lines=['# Fase1/H3 — auditoría de llaves del libro original','',f'B/C — Fuente `{SRC.name}`; SHA-256 `{before}`. Lectura de {len(a):,} filas de hoja9 y {len(m):,} filas de hoja11. No se modifica ningún archivo fuente ni denominador.','', '| Control observado B/C | Filas de hoja9 |','|---|---:|']
for k,v in summary.items():lines.append(f'| {k} | {v} |')
lines+=['','## Por cohorte','', '| Cohorte | Filas | Retenidos | Llave repetida | Varios códigos candidatos | Código no coincide con primera fila MU | Código fuera de todos los candidatos |','|---|---:|---:|---:|---:|---:|---:|']
for year,c in sorted(byyear.items()):lines.append(f"| {year} | {c['filas_analiticas']} | {c['retenidos']} | {c['llave_repetida']} | {c['varios_codigos_completos']} | {c['codigo_excel_no_coincide_primera_fila_mu']} | {c['codigo_excel_no_en_ninguna_fila_mu_de_llave']} |")
lines+=['','C — U obtiene COD_CAR por VLOOKUP del nombre; V/W usan primera coincidencia de cohorte+N_DOC+DV+COD_CAR; X concatena sede/version de la primera fila MU y jornada de hoja9. La clave no incluye tipo documental, jornada, modalidad ni versión.','',f"C — QA: {json.dumps(q,ensure_ascii=False)}.",'','C/E — Una llave repetida no demuestra un error: separar filas iguales, mismo código y códigos distintos. Coincidencia de componentes tampoco acredita por sí sola CODCLI/plan ni aceptación SIES. Los candidatos se preservan para revisar; no se adjudica otra versión.','', 'C — Detalle con claves/fila/candidatos solo en ruta privada registrada en RESULTADOS_AGREGADOS.json. Las tablas de este informe son agregadas.','', 'E — Este hito no amplía a Fase2 ni declara resuelta la equivalencia institucional completa.']
(O/'INFORME_LLAVES.md').write_text('\n'.join(lines)+'\n')
print(json.dumps({'conteos':dict(summary),'QA':q,'duplicados_mu_completa':report['duplicados_mu_completa'],'duplicados_llaves_usadas':report['duplicados_llaves_usadas'],'por_cohorte':dict(byyear)},ensure_ascii=False))
