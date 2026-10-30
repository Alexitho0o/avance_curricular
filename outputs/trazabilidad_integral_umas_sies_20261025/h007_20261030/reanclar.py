"""H007: consulta de excepciones congeladas y nuevos agregados, sin adjudicar.

No recalcula tasas gobernadas, enlaces ni el diccionario. MRUN se agrupa en
memoria separada; nunca se combina con identidades UMAS. Salidas solo agregadas.
"""
from pathlib import Path
from collections import Counter, defaultdict
import csv, hashlib, json, re, sys
import openpyxl
sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).parent
PRIV = Path.home()/'.local/share/ipss/h007_20261030'
P = PRIV.parent
sys.path.insert(0, str(ROOT))
from common.diccionario_sies_seguimiento import DiccionarioSiesSeguimiento, rut_normalizado
SOURCES = {}

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def source(p, expected=None):
    p = Path(p)
    if not p.is_absolute(): p = ROOT/p
    h = sha(p)
    if expected: assert h == expected, f'Hash distinto: {p.name}'
    SOURCES[str(p)] = h
    return p

def load(p, expected=None): return json.loads(source(p, expected).read_text())
def txt(v): return '' if v is None else str(int(v)) if isinstance(v,float) and v.is_integer() else str(v).strip()
def write(name, d):
    (OUT/name).write_text(json.dumps(d, ensure_ascii=False, indent=2, sort_keys=True)+'\n')
def counter_rows(c, fields):
    return [dict(zip(fields,k),n=v) for k,v in sorted(c.items(),key=lambda x:tuple(map(str,x[0])))]
def table(rows, cols):
    return '\n'.join(['| '+' | '.join(cols)+' |','|'+'|'.join(['---']*len(cols))+'|']+['| '+' | '.join(str(r.get(k,'')) for k in cols)+' |' for r in rows])

def build():
    source('common/analisis_desercion/h001.py')
    source('outputs/paridad_cohorte_2007_2026/construir_cohortes.py')
    slots = load(P/'trazabilidad_integral_fase3_20261025/VISTA_T1_T2_T3.json','72edecb9507f46985533d1c63b83735beb4db58d743c89c226ca6be1a2c8ef31')
    source(Path.home()/'Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261025_V24.xlsx','bfeba67628107f460cad69c5db119a6245782142bd5b65ab435b0cd116bca2db')
    # Valores de una copia recalculada ya gobernada en R19, sin recalcular V24.
    wb = openpyxl.load_workbook(source(P/'fase_h_regla_intertrimestral_20261027/V24_QA_COPIA.xlsx','1e28d5955d874e7d86ec0ca87ab6a2cc478515c0dbd604f4fc86f60fa491c481'),read_only=True,data_only=True)
    photo = dict(enumerate(wb['20_Foto_BBDD'].iter_rows(values_only=True),1))
    active = load('data/diccionario_sies_seguimiento/activo.json')
    desc = load(active['descriptor'],active['sha256'])
    source(desc['base_privada']['ruta'],desc['base_privada']['sha256'])
    api = DiccionarioSiesSeguimiento(root=ROOT)
    idx = defaultdict(list)
    for (payload,) in api.db.execute('SELECT payload FROM observaciones'):
        r = json.loads(payload);idx[(r['persona'],txt(r['ano']),txt(r['periodo']))].append(r)
    api.close()
    absent = [s for s in slots if not s['contextos']]
    assert len(slots)==6358 and len(absent)==3520
    classification=Counter(); details=[]; photo_counts=Counter()
    # El vínculo fila_foto ya existe: se verifica documento y año, no se crea otro.
    for s in absent:
        r=photo[s['fila_origen_foto']]
        assert txt(r[1])==txt(s['anio_origen']) and txt(r[4])==txt(s['documento_original'])
        period=s['periodo_nativo']; state=txt(r[12] if period=='2' else r[15])
        rr=idx[(s['persona'],txt(s['anio_observado']),period)]
        if rr:
            cause='CONTEXTO_PRESENTE_EN_V20_AUSENTE_EN_INSTANTANEA_V8'
            requirement='Integrar una nueva vista con autorización; V24 conserva v8. Contexto de la misma persona/año/período, no prueba de mismo expediente T1 ni vigencia.'
            cls='B/C'
        elif state=='VIG 1':
            cause='FOTO_INGRESO_VIG1_SIN_CONTEXTO_DE_ACTIVIDAD';cls='B/E'
            requirement='Actividad/expediente del período con identidad y calendario acreditados; VIG 1 no aporta un contexto PROMEDIOS.'
        elif state=='VIG 0':
            cause='FOTO_INGRESO_VIG0_SIN_CONTEXTO_DE_ACTIVIDAD';cls='B/E'
            requirement='Causa académica/administrativa y cobertura del período; VIG 0 observado no prueba abandono ni explica por sí solo la ausencia de actividad.'
        elif state=='Sin matrícula de ingreso en el corte':
            cause='MATRICULA_DE_INGRESO_NO_LOCALIZADA_EN_FOTO';cls='B/E'
            requirement='Expediente de ingreso en corte y revisión de selección/llave; no significa que la persona carezca de otra matrícula.'
        elif state=='NO APLICA':
            cause='NO_APLICA_EN_FOTO_SIN_CONTEXTO';cls='B/E'
            requirement='Conservar exclusión temporal de la foto; no retirar la posición de V24 ni certificar ausencia de actividad.'
        else:
            cause='FOTO_NO_COMPLETA_Y_SIN_CONTEXTO';cls='E'
            requirement='Foto completa del mismo corte y correspondencia del período nativo con el trimestre; no clasificable como no matriculado ni falla de llave.'
        key=(s['anio_origen'],s['anio_observado'],period,cause,cls,requirement)
        classification[key]+=1;photo_counts[(s['anio_origen'],s['anio_observado'],period,state)]+=1
        details.append({'fila_vista_origen':s['fila_origen_foto'],'periodo':period,'causa':cause,'estado_foto':state,'observaciones_v20':len(rr),'clase':cls})
    # Detalle sin documentos, pero por fila individual: permanece privado.
    PRIV.mkdir(exist_ok=True,mode=0o700)
    (PRIV/'POSICIONES_3520.json').write_text(json.dumps(details,ensure_ascii=False,indent=2));(PRIV/'POSICIONES_3520.json').chmod(0o600)
    absent_rows=counter_rows(classification,['cohorte_T1','anio_observado','periodo_nativo','causa','clase','requisito'])
    write('POSICIONES_SIN_CONTEXTO.json',{'unidad':'posiciones origen × período; no personas','total':3520,'clasificacion':absent_rows,'foto_observada':counter_rows(photo_counts,['cohorte_T1','anio_observado','periodo_nativo','estado_foto']),'limite':'T2/T3 son etiquetas de seguimiento; período nativo no se homologa a trimestre/fecha sin calendario. Causas disjuntas con prioridad contexto v20. Ausencia de contexto no acredita ausencia de matrícula.'})

    parity=load('data/paridad_codigo_sies/20261029/TABLA_PARIDAD_2026_2027.json')['filas']
    profiles=list(csv.DictReader(source('data/paridad_codigo_sies/regla_operativa_mu_v2/TIPO_PLAN_POR_PLAN_UMAS.csv').open(),delimiter=';'))
    coherence=load('outputs/excel_resumen_cruce_20261030/COHERENCIA_TIPO_H005C.json')
    pending=[c for s in slots for c in s['contextos'] if not c['codigo_destino']]
    assert len(pending)==99
    pc=Counter(); categories=Counter(); private=[]
    for c in pending:
        old='JORNADA_DISTINTA' if c['dimensiones_contexto_distintas_origen'] else 'SIN_ENLACE_EXPEDIENTE_T1'
        types={p['tipo_plan'] for p in profiles if p['codcarpr']==c['carrera_umas'] and p['plan_umas']==c['plan_destino']}
        options=[r for r in parity if r['anio']==2026 and c['carrera_umas'] in r['codcarpr'] and c['plan_destino'] in r['plan_umas'] and txt(r['cod_sede'])==txt(c['sede_contexto']) and txt(r['cod_jornada'])==txt(c['jornada_contexto'])]
        options=[r for r in options if ('REGULAR' if r['tipo_plan_codigo']=='1' else 'CONTINUIDAD' if r['tipo_plan_codigo']=='3' else 'E') in types] if types and types!={'E'} else options
        codes=sorted({r['codigo_unico'] for r in options})
        status='CATALOGO_CANDIDATO_UNICO_NO_ATRIBUCION' if len(codes)==1 else 'CATALOGO_MULTIPLE_E' if codes else 'SIN_CANDIDATO_CATALOGO_E'
        # Incluso único: H002 2026/2027 y tipo no acreditan atribución al período 2025/26.
        requirement='Código documental del expediente/período o decisión D explícita; no heredar jornada distinta.' if old=='JORNADA_DISTINTA' else 'Enlace del expediente T1 al código, con contexto temporal; catálogo no prueba identidad del expediente.'
        typestr=','.join(sorted(types)) if types else 'SIN_PERFIL'
        inherited_coherence=[x for x in coherence['cruce_tipos'] if x.get('anio')==c['anio_observado'] and x.get('tipo_umas') in types]
        pc[(c['anio_origen'],c['anio_observado'],c['periodo_nativo'],old,c['carrera_umas'],c['plan_destino'],txt(c['sede_contexto']),txt(c['jornada_contexto']),c['codigo_origen_T1'],typestr,','.join(codes),status,requirement)]+=1
        categories[(old,status)]+=1
        private.append({'observacion_id':c['observacion_id'],'fila_origen':c['fila_origen_foto'],'periodo':c['periodo_nativo'],'estado':status,'candidatos_catalogo':codes,'requisito':requirement,'respaldado_personal':False})
    (PRIV/'PENDIENTES_99.json').write_text(json.dumps(private,ensure_ascii=False,indent=2));(PRIV/'PENDIENTES_99.json').chmod(0o600)
    fields=['cohorte_T1','anio_observado','periodo_nativo','motivo_v24','carrera_umas','plan','sede','jornada','codigo_T1','tipo_v2','candidatos_H002_2026','estado','requisito']
    write('REEVALUACION_99.json',{'clase':'C/E; catálogo H002 con decisiones D delimitadas','total':99,'pasan_a_respaldados_personales':0,'siguen_pendientes_personales':99,'nueva_herencia':False,'resumen':counter_rows(categories,['motivo_v24','estado']),'casos_agrupados':counter_rows(pc,fields),'coherencia_reutilizada':coherence['anual'],'limite':'La coherencia es agregada por año; no acredita un código individual. No se consultó v20 para adjudicar estos 99: solo H002, tipo v2 y coherencia, sobre sus contextos congelados. H002 2026/2027 no se retrotrae a 2025.'})

    # Hito 3: agregación separada, sin llave común entre los sistemas.
    catalog=load('data/catalogo_fuentes/CATALOGO_FUENTES.json')['catalogo']
    mc=defaultdict(set); rawcounts=Counter();level={}
    for filename in ['reporte_matricula_ip_san_sebastian.csv','MRUN_MATRICULA_IPSS_162_2026_H2.csv']:
        e=next(e for e in catalog if e['archivo']==filename)
        path=source(e['ruta_consolidada'],e['sha256'])
        for r in csv.DictReader(path.open(encoding='utf-8-sig')):
            y=int(r['cat_periodo'])
            if y not in range(2022,2027):continue
            assert r['mrun'].strip(), 'Registro SIES sin identificador agrupable'
            key=(y,r['codigo_unico']);mc[key].add(r['mrun']);rawcounts[key]+=1;level[key]=r['nivel_global']
    # Reutiliza selección ya congelada de H002g; no recalcula enlace ni VIG.
    links=load(P/'h002g_20261029/ENLACES_MU_DA.json')
    mu=defaultdict(set);muall=defaultdict(set)
    for r in links:
        key=(r['anio'],r['codigo_sies']);identity=tuple(r['identidad']);muall[key].add(identity)
        if r['incluida']:mu[key].add(identity)
    aggregate=[]
    for key in sorted(set(mc)|set(muall)):
        y,c=key;has_s=key in mc;has_mu=key in muall
        diff=len(mu[key])-len(mc[key]) if has_s and has_mu else None
        aggregate.append({'anio':y,'codigo_sies':c,'nivel_sies':level.get(key,'SIN_PAR_SIES'),'referencia_personas_MRUN':len(mc[key]) if has_s else None,'observado_personas_MU_seleccionada':len(mu[key]) if has_mu else None,'personas_MU_todos_VIG':len(muall[key]) if has_mu else None,'registros_sies':rawcounts[key] if has_s else None,'diferencia_observado_menos_referencia':diff,'estado':'SIN_PAR' if diff is None else 'CERO' if diff==0 else 'DISTINTA_DE_CERO','clase':'B/C; diferencia no explica causa individual'})
    annual=[]
    for y in range(2022,2027):
        rr=[x for x in aggregate if x['anio']==y];states=Counter(x['estado'] for x in rr)
        annual.append({'anio':y,'referencia_personas_codigo':sum(x['referencia_personas_MRUN'] or 0 for x in rr),'observado_personas_codigo':sum(x['observado_personas_MU_seleccionada'] or 0 for x in rr),'diferencia_totales':sum(x['observado_personas_MU_seleccionada'] or 0 for x in rr)-sum(x['referencia_personas_MRUN'] or 0 for x in rr),'codigos_cero':states['CERO'],'codigos_diferencia':states['DISTINTA_DE_CERO'],'codigos_sin_par':states['SIN_PAR'],'registros_sies_menos_personas_codigo':sum((x['registros_sies'] or 0)-(x['referencia_personas_MRUN'] or 0) for x in rr)})
    by_level=[]
    for y in range(2022,2027):
        for lev in sorted({x['nivel_sies'] for x in aggregate if x['anio']==y}):
            rr=[x for x in aggregate if x['anio']==y and x['nivel_sies']==lev]
            cc=Counter(x['estado'] for x in rr)
            by_level.append({'anio':y,'nivel':lev,'referencia':sum(x['referencia_personas_MRUN'] or 0 for x in rr),'observado':sum(x['observado_personas_MU_seleccionada'] or 0 for x in rr) if any(x['observado_personas_MU_seleccionada'] is not None for x in rr) else None,'cero':cc['CERO'],'distinta':cc['DISTINTA_DE_CERO'],'sin_par':cc['SIN_PAR']})
    write('MRUN_AGREGADO.json',{'clase':'B/C; causas E','unidad':'personas distintas dentro de cada código/año; sumas son incidencias persona-código, no personas institucionales únicas','foto_referencia_interna':'MU anual 142342, selección congelada H002g; no padrón restringido T1 de V24 ni DA actual','regla_seleccion':'Reutiliza incluida de ENLACES_MU_DA; 2025 incluye VIG0; los demás años selección anterior H002g, sin cambios. Todos VIG se muestran aparte. Postítulos permanecen separados por nivel; sin par MU no es cero observado.','sin_enlace_identidad':True,'reutilizacion':'R10 controla filas por código/año; no contiene MRUN distintos. Este conteo nuevo no repite la paridad de filas ni reabre la brecha R09.','anual':annual,'por_nivel':by_level,'por_codigo':aggregate})

    # Hito 4: intersección de DOS resultados existentes; no recalcular H001.
    errors=[e for e in coherence['errores'] if e['anio']==2024 and e['tipo_umas']=='CONTINUIDAD' and e['tipo_sies']=='REGULAR']
    target={(e['codcarpr'],e['plan'],e['codigo']):e['matriculas'] for e in errors}
    selected=[]
    for r in links:
        if r['anio']!=2024 or not r['incluida'] or not r['seleccion']:continue
        key=(r['seleccion'][1],r['plan'],r['codigo_sies'])
        if key in target:selected.append((key,r))
    assert Counter(k for k,r in selected)==Counter(target), 'El resultado privado no reproduce los 131 gobernados'
    eligible={};total=retained=0
    for row in wb['9_Base_Analitica_Detalle'].iter_rows(min_row=7,values_only=True):
        if txt(row[1])!='2024':continue
        total+=1;retained+=int(row[5]);key=(txt(row[6]),txt(row[7]),txt(row[8]).upper())
        if key[0]=='R':
            body,dv=rut_normalizado(key[1]+'-'+key[2]).split('-');key=(key[0],body,dv)
        assert key not in eligible;eligible[key]=int(row[5])
    assert (total,retained)==(558,326),'La foto no reproduce H001 2024, detener escenario'
    mu2024={tuple(r['identidad']) for r in links if r['anio']==2024}
    assert set(eligible).issubset(mu2024), 'La llave no valida los 558 H001 contra el resultado MU existente'
    assert len({tuple(r['identidad']) for k,r in selected})==131, '131 matrículas no equivalen a 131 personas'
    intersect=[];h4c=Counter()
    for k,r in selected:
        identity=tuple(txt(x) for x in r['identidad'])
        yes=identity in eligible
        h4c[k+('EN_H001' if yes else 'FUERA_H001',eligible.get(identity,'NO_APLICA'))]+=1
        if yes:intersect.append((identity,eligible[identity]))
    assert len(set(i for i,_ in intersect))==len(intersect)
    origin_counts=Counter((k[0],r['cohorte_origen_sies'][0],txt(r['cohorte_origen_sies'][1])) for k,r in selected)
    n=len(intersect);ret=sum(v for _,v in intersect)
    scenario=((326-ret)/(558-n)-326/558)*100 if n<558 else None
    write('ESCENARIO_131_H001.json',{'clase':'C escenario; causalidad E','evidencia':'COHERENCIA_TIPO_H005C + ENLACES_MU_DA congelados + V24_QA_COPIA hoja 9; no reproducción de H001 ni nuevo enlace UMAS','continuidad_en_regular_2024':131,'en_cohorte_H001':n,'retenidos_en_interseccion':ret,'fuera_cohorte_H001':131-n,'H001_preservado':{'casos':558,'retenidos':326,'personas_h001_localizadas_en_MU_2024':len(set(eligible)&mu2024),'personas_distintas_en_los_131':131},'escenario_exclusion_interseccion':{'casos':558-n,'retenidos':326-ret,'cambio_pp':scenario,'formula':'100*((326-r)/(558-n)-326/558)','condicion':'Supuesto de excluir del denominador y numerador los planes de continuidad; no aprobado ni aplicado'},'cohorte_origen_131':counter_rows(origin_counts,['carrera','anio_origen','semestre_origen']),'pp_explicados_por_inclusion_en_H001':0.0 if n==0 else None,'pp_causalmente_explicados':None,'requisito_causal':'No hay cohorte nominal SIES oficial 2024 enlazable autorizada ni contrafactual de retención; cambiar versión no prueba cambio en permanencia institucional. La sensibilidad no identifica la causa de -9,08 pp.','detalle_agrupado':counter_rows(h4c,['carrera','plan','codigo','pertenencia','retenido'])})
    wb.close()
    write('FUENTES_H007.json',{'fuentes':[{'ruta':p,'sha256':h} for p,h in sorted(SOURCES.items())],'clase':'B fuentes / C consultas; ninguna fuente modificada'})
    print(json.dumps({'3520':dict(Counter({cause:sum(x['n'] for x in absent_rows if x['causa']==cause) for cause in {x['causa'] for x in absent_rows}})),'99':counter_rows(categories,['motivo','estado']),'MRUN':annual,'H4':{'en_H001':n,'retenidos':ret,'escenario_pp':scenario}},ensure_ascii=False))

if __name__=='__main__':build()
