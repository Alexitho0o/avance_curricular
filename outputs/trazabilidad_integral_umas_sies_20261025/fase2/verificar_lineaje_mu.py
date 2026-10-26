"""Auditoría de lectura: originales anuales -> BASE_RETENCION_MU.

No exporta identificadores ni modifica entradas. Emite controles agregados.
No asigna versiones: compara las ya escritas en ambas fuentes.
"""
from pathlib import Path
import csv
import json
import hashlib
from collections import Counter
from decimal import Decimal, InvalidOperation
from datetime import datetime
from zoneinfo import ZoneInfo
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
BASE_DIR = ROOT / "outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342"
MU = BASE_DIR / "BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx"
FIELDS = ["TIPO_DOC", "N_DOC", "DV", "COD_SED", "COD_CAR", "JOR", "VERSION", "VIG"]
NUMERIC = {"N_DOC", "COD_SED", "COD_CAR", "JOR", "VERSION", "VIG", "COD_IES"}
SHEETS = {2025: "08-05-2025_22-26-22_Matrícula_P", 2026: "CONSOLIDADO"}

def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024*1024), b""):
            h.update(chunk)
    return h.hexdigest()

def norm(field, value):
    text = "" if value is None else str(value).strip()
    # Transformación ya implementada por el constructor anual; no regla nueva.
    if field == "VIG" and text.upper() in {"VIGENTE", "NO VIGENTE"}:
        return {"VIGENTE": "1", "NO VIGENTE": "0"}[text.upper()]
    if field in NUMERIC and text:
        try:
            d = Decimal(text)
            return str(int(d)) if d == d.to_integral() else str(d)
        except InvalidOperation:
            pass
    return text

def read(path, year):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="latin-1" if year == 2024 else "utf-8-sig", newline="") as f:
            rows = list(csv.reader(f, delimiter=";"))
        headers = rows[0]
        source = [(i, r) for i, r in enumerate(rows[1:], 2) if any(str(x).strip() for x in r)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[SHEETS[year]]
        it = ws.iter_rows(values_only=True)
        headers = next(it)
        source = [(i, r) for i, r in enumerate(it, 2) if any(x is not None and str(x).strip() for x in r)]
        wb.close()
    headers = [str(h).strip().upper() for h in headers]
    assert all(headers.count("VIGENCIA" if year == 2025 and f == "VIG" else f) == 1 for f in FIELDS)
    positions = [headers.index("VIGENCIA" if year == 2025 and f == "VIG" else f) for f in FIELDS]
    data = [(i, tuple(norm(f, r[j]) for f, j in zip(FIELDS, positions))) for i, r in source]
    institution = Counter(norm("COD_IES", r[headers.index("COD_IES")]) for _, r in source) if "COD_IES" in headers else None
    return data, institution

def main():
    inventory_path = BASE_DIR / "INVENTARIO_FUENTES_SHA256.json"
    inventory = json.loads(inventory_path.read_text())
    before = {str(MU): sha(MU), str(inventory_path): sha(inventory_path)}
    wb = openpyxl.load_workbook(MU, read_only=True, data_only=True)
    ws = wb["BASE_RETENCION_MU"]
    iterator = ws.iter_rows(values_only=True)
    headers = list(next(iterator))
    observed = {}
    for rownum, row in enumerate(iterator, 2):
        year = int(row[headers.index("ANIO_INFORMADO")])
        values = tuple(norm(f, row[headers.index(f)]) for f in FIELDS)
        observed.setdefault(year, []).append((rownum, values))
    wb.close()
    checks = []
    for spec in inventory:
        year, path = spec["anio"], Path(spec["archivo"])
        before[str(path)] = sha(path)
        raw, institutions = read(path, year)
        current = observed[year]
        raw_values = [v for _, v in raw]
        current_values = [v for _, v in current]
        missing = sum((Counter(raw_values) - Counter(current_values)).values())
        extra = sum((Counter(current_values) - Counter(raw_values)).values())
        different = {f: sum(a[j] != b[j] for a, b in zip(raw_values, current_values)) for j, f in enumerate(FIELDS)}
        checks.append({
            "clasificacion": "B/C", "anio": year, "fuente": str(path),
            "hoja": SHEETS.get(year, "CSV con encabezado"),
            "sha256": before[str(path)], "hash_coincide_inventario": before[str(path)] == spec["sha256"],
            "filas_fuente": len(raw), "filas_base": len(current),
            "rango_filas_fuente": [raw[0][0], raw[-1][0]],
            "rango_filas_base": [current[0][0], current[-1][0]],
            "columnas_comparadas": FIELDS, "diferencias_por_columna_mismo_orden": different,
            "faltantes_multiconjunto": missing, "adicionales_multiconjunto": extra,
            "mismo_orden_y_valores": raw_values == current_values,
            "COD_IES_observados": dict(institutions) if institutions is not None else "No existe columna: alcance IPSS documentado en inventario",
            "estado": "PASA" if raw_values == current_values and before[str(path)] == spec["sha256"] else "FALLA",
        })
    preserved = [{"ruta": p, "sha256": h, "sin_cambios": sha(Path(p)) == h} for p, h in before.items()]
    result = {
        "clasificacion": "B/C", "fecha": datetime.now(ZoneInfo("America/Santiago")).isoformat(),
        "metodo": "Comparación independiente en memoria de identidad, sede, carrera, jornada, versión y vigencia, por año, como multiconjunto y en orden original. Números enteros se normalizan solo para comparar; entradas intactas. VIGENCIA 2025: Vigente -> 1 y No vigente -> 0, transformación existente en scripts/construir_base_retencion_mu_2022_2026.py:typed_value.",
        "incidencia_del_lector_resuelta_c": "La primera prueba detectó 2371 diferencias de representación VIGENCIA textual frente a VIG numérica en 2025. Relectura: Vigente=2298 y No vigente=73. Se incorporó al lector de auditoría la transformación existente del constructor y se repitió el control; no se modificó ninguna base.",
        "base": str(MU), "hoja_base": "BASE_RETENCION_MU", "sha256_base": before[str(MU)],
        "filas_totales": sum(c["filas_base"] for c in checks), "controles": checks,
        "preservacion": preserved,
        "limites_e": ["No acredita recepción/aceptación SIES.", "No verifica CODCLI ni asigna una versión nueva.", "No verifica las demás 24 columnas del esquema rector ni recalcula elegibilidad/retención."],
        "qa": "PASA" if all(c["estado"] == "PASA" for c in checks) and all(p["sin_cambios"] for p in preserved) else "FALLA"
    }
    (OUT / "QA_LINEAJE_MU.json").write_text(json.dumps(result, ensure_ascii=False, indent=2)+"\n")
    print(json.dumps({"qa": result["qa"], "filas": result["filas_totales"], "por_anio": [{"anio": c["anio"], "filas": c["filas_base"], "estado": c["estado"]} for c in checks]}, ensure_ascii=False))

if __name__ == "__main__":
    main()
