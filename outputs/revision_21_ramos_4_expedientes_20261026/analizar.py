"""Analiza casos ambiguos en memoria y escribe detalle exclusivamente privado."""
import collections,csv,hashlib,json,os,re,sqlite3,subprocess
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2]; OUT=Path(__file__).resolve().parent
previous=json.loads((ROOT/'outputs/conciliacion_22_ramos_titulados_20261026/EVIDENCIA_AGREGADA.json').read_text())
source=previous['fuentes'];book=Path(source['respaldo_xlsx']['ruta']);prom=Path(source['promedios']['ruta']);db=Path('/Users/alexi/.local/share/ipss/v11_fase2_20261026/DICCIONARIO_SIES_CANDIDATO_V11_86_Y_AMBIGUEDAD.sqlite')
data_dir=Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/IPSS_CONCILIACION_UMAS_SIES_BASES/03_UMAS_DATOSDEALUMNOS')
def sha(path):
 h=hashlib.sha256()
 with path.open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
assert sha(book)==source['respaldo_xlsx']['sha256'] and sha(prom)==source['promedios']['sha256'] and sha(db)==source['v11_sha256']
def s(v):
 if v is None:return ''
 if isinstance(v,(int,float)) and v==int(v):return str(int(v))
 return str(v).strip().upper()
def rut(n,d):
 a=s(n).replace('.','').replace(' ','');b=s(d)
 return str(int(a))+'-'+b if a.isdigit() and re.fullmatch('[0-9K]',b) else None
targets={(x['carrera'],x['plan']) for x in previous['titulados_por_plan']}
conn=sqlite3.connect(db.as_uri()+'?mode=ro',uri=True);idx=collections.defaultdict(set); obs=collections.defaultdict(list)
for person,codcli,career,plan,year,period,sede,jor,payload in conn.execute('select persona,codcli,carrera,plan,anio,periodo,sede,jornada,payload from observaciones'):
 o=json.loads(payload);code=o.get('codigo_documental_atribuible') or o.get('codigo_interno_adoptado')
 if o.get('tipo_doc_normalizado')=='RUT' and code:
  idx[(person,code)].add((codcli,career,plan))
  obs[person].append({'CODCLI':codcli,'carrera':career,'plan':plan,'ano':year,'periodo':period,'sede':sede,'jornada':jor,'codigo':code,'tipo':'documental' if o.get('codigo_documental_atribuible') else 'interno'})
conn.close(); w=openpyxl.load_workbook(book,read_only=True,data_only=True)
def getrows(sheet):
 it=w[sheet].iter_rows(values_only=True);head=[s(v) for v in next(it)]
 for n,row in enumerate(it,2):yield n,dict(zip(head,row))
titles=collections.defaultdict(list)
for sheet in ('Res_Tit_Carga','5319 Titulados Histór'):
 for rn,r in getrows(sheet):
  if s(r.get('TIPO_DOCUMENTO'))!='R':continue
  person=rut(r.get('NUM_DOCUMENTO'),r.get('DV'));code=s(r.get('COD_SIES_OBT_TIT') or r.get('COD_UNICO_OBT_TIT'));year=s(r.get('FECHA_OBT_TIT_GRA') or r.get('AÑO_TITULACION'))
  if person and code and any((career,plan) in targets for _,career,plan in idx[(person,code)]) and len(idx[(person,code)])>1:
   titles[person].append({'hoja':sheet,'fila':rn,'codigo':code,'fecha_o_ano':year,'expedientes_compatibles':[{'CODCLI':c,'carrera':ca,'plan':pl} for c,ca,pl in sorted(idx[(person,code)])]})
trace=collections.defaultdict(list)
for rn,r in getrows('Res_Tit_Unificada'):
 person=rut(r.get('NUM_DOCUMENTO'),r.get('DV'))
 if person in titles:
  trace[person].append({'hoja':'Res_Tit_Unificada','fila':rn,'codigo':s(r.get('COD_SIES_OBT_TIT')),'fecha_o_ano':s(r.get('FECHA_OBT_TIT_GRA')),'codcli_candidato':s(r.get('TRACE_CODCLI_DATOS_ALUMNOS')),'carrera_candidata':s(r.get('TRACE_CODCARPR_DATOS_ALUMNOS')),'puntaje':s(r.get('TRACE_SCORE_CANDIDATO')),'razones':s(r.get('TRACE_REASONS_CANDIDATO')),'observacion':s(r.get('TRACE_OBSERVACION'))})
w.close()
records=collections.defaultdict(list); sources=[]
for f in (data_dir/'DATOSDEALUMNOS_12159.xlsx',data_dir/'DATOSDEALUMNOS_11248.xlsx'):
 hsh=sha(f);dw=openpyxl.load_workbook(f,read_only=True,data_only=True);it=dw['Hoja1'].iter_rows(values_only=True);head={s(v):i for i,v in enumerate(next(it)) if v is not None}
 for rn,row in enumerate(it,2):
  val=lambda k:s(row[head[k]]) if k in head else ''
  raw=val('RUT').replace('.','').replace(' ','');m=re.fullmatch(r'(\d+)-?([0-9K])',raw);person=rut(m.group(1),m.group(2)) if m else None
  if person in titles:records[person].append({'fuente':f.name,'fila':rn,'CODCLI':val('CODCLI'),'carrera':val('CODCARPR'),'ano_ingreso':val('ANOINGRESO'),'periodo_ingreso':val('PERIODOINGRESO'),'ano_matricula':val('ANOMATRICULA'),'periodo_matricula':val('PERIODOMATRICULA'),'fecha_matricula':val('FECHAMATRICULA'),'situacion':val('SITUACION'),'estado':val('ESTADOACADEMICO'),'jornada':val('JORNADA'),'sede':val('SEDE')})
 dw.close();assert sha(f)==hsh;sources.append({'archivo':f.name,'sha256':hsh})
pw=openpyxl.load_workbook(prom,read_only=True,data_only=True);it=pw['Hoja1'].iter_rows(values_only=True);head={s(v):i for i,v in enumerate(next(it)) if v is not None};activity=collections.defaultdict(lambda:collections.defaultdict(set))
for row in it:
 get=lambda k:s(row[head[k]]) if k in head else ''
 person=rut(get('RUT'),get('DIG'))
 if person in titles:activity[person][(get('CODCLI'),get('CODCARR'),get('PLAN_DE_ESTUDIO'))].add((get('ANO'),get('PERIODO')))
pw.close();assert sha(prom)==source['promedios']['sha256']
cases=[];private=[]
for i,person in enumerate(sorted(titles),1):
 tr=titles[person]; pairs={(e['carrera'],e['plan']) for t in tr for e in t['expedientes_compatibles']};periods=[{'carrera':ca,'plan':pl,'periodos':sorted([{'ano':y,'periodo':p} for y,p in dates],key=lambda x:(x['ano'],x['periodo']))} for (cl,ca,pl),dates in sorted(activity[person].items())]
 matches=[x for x in trace[person] if any(x['codigo']==y['codigo'] and x['fecha_o_ano']==y['fecha_o_ano'] for y in tr)]
 cases.append({'caso':f'C{i:02}','planes_compatibles':[{'carrera':a,'plan':b} for a,b in sorted(pairs)],'actividad_por_expediente_sin_identificador':periods,'titulos':[{k:v for k,v in t.items() if k!='expedientes_compatibles'} for t in tr],'trazas_titulacion_coincidentes':len(matches),'estado':'E: confirmar plan individual, fecha efectiva y validez de traza calculada'})
 private.append({'caso':f'C{i:02}','RUT':person,'titulos':tr,'trazas_calculadas':trace[person],'datosalumnos':records[person],'contextos_v11':obs[person]})
assert len(cases)==4, f'Se observaron {len(cases)} personas, distinta cantidad de lo informado; no continuar automáticamente.'
fdir=Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026');fdir.mkdir(mode=0o700,parents=True,exist_ok=True)
assert subprocess.run(['git','-C',str(fdir),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
priv=fdir/'DETALLE_CUATRO_CASOS_V2.json';assert not priv.exists()
with os.fdopen(os.open(priv,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600),'w') as f:json.dump(private,f,ensure_ascii=False,indent=2)
# Estructura conserva equivalencias literales; comparar ambas direcciones sin transitividad.
structure=json.loads(Path(source['estructura']['ruta']).read_text());edge=[]
for x in structure:
 for e in x.get('equivalencias_detalle',[]):edge.append(((s(x['CODPESTUD']),s(x['CODRAMO'])),(s(e['plan_destino']),s(e['ramo_destino'])),e.get('filas_fuente',[])))
res=[]
for item in previous['discrepancias']:
 if item['relaciones_catalogo']:continue
 dst=(item['plan'],item['asignatura']);candidates=set(item['equivalente_exacto_en_plan']+item['mismo_nombre_en_plan']);found=[]
 for a,b,lines in edge:
  if (a==dst and b[0]==dst[0] and b[1] in candidates) or (b==dst and a[0]==dst[0] and a[1] in candidates):found.append({'desde':list(a),'hacia':list(b),'filas_fuente':lines})
 res.append({'carrera':item['carrera'],'plan':item['plan'],'ramo':item['asignatura'],'candidatos_misma_malla':sorted(candidates),'relacion_explicita_directa':found,'estado':'B/C relación observada; E aplicación y fechas' if found else 'E: no hay relación explícita para candidatos en esta fuente'})
assert len(res)==21
case_groups=collections.Counter(tuple((p['carrera'],p['plan']) for p in c['planes_compatibles']) for c in cases)
grouped_cases=[{'planes_compatibles':[{'carrera':a,'plan':b} for a,b in key],'personas':count,'estado':'E: confirmar atribución y fecha efectiva'} for key,count in sorted(case_groups.items())]
result={'clasificacion':'B registros; C cruce exacto; E atribución y vigencia','fecha':'2026-10-26','fuentes':{'respaldo_titulacion':source['respaldo_xlsx'],'promedios':source['promedios'],'estructura':source['estructura'],'datos_alumnos':sources,'v11_sha256':sha(db)},'ramos':res,'casos_agrupados':grouped_cases,'privado':{'ruta':str(priv),'sha256':sha(priv)},'limitaciones':'Una traza de candidatura no es título con plan declarado. La actividad no fecha migración ni cierre; fuentes con corte no certificado. No se cambian bases ni denominadores.'}
(OUT/'RESULTADO_AGREGADO.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'ramos':len(res),'relaciones_explicitas':sum(bool(x['relacion_explicita_directa']) for x in res),'personas':len(cases),'trazas_coincidentes':sum(x['trazas_titulacion_coincidentes'] for x in cases),'privado_sha256':sha(priv)},ensure_ascii=False))
