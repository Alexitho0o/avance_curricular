"""Comprueba cifras publicadas IPSS y contrasta con variantes agregadas existentes."""
import csv
import hashlib
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BASE = Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/0 BASES DE DATOS/RETENCION')
PUBLICACIONES = [
    ('Informe_Retencion_SIES_2025.xlsx', 'e5b21fd4194725ed57a21c7e6511aa90c58cdab298d97a51ccf430861638f82d', 44, 63),
    ('Informe_Retencion_SIES_2026_270702026.xlsx', '03897a1604aa656f9ff60e232eda739e29f4912123e5d7127d6551a37957a8a1', 40, 59),
]
for name, expected, header_row, row_number in PUBLICACIONES:
    p = BASE / name
    with p.open('rb') as f:
        assert hashlib.file_digest(f, 'sha256').hexdigest() == expected
    sheet = load_workbook(p, read_only=True, data_only=True)['Retención 1er año x IES']
    # Find the exact institutional header within the preceding block.
    header_row = next(r for r in range(row_number - 1, 1, -1)
                      if sheet.cell(r, 1).value == 'Cod. Institución')
    assert sheet.cell(row_number, 1).value == 162
    assert sheet.cell(row_number, 2).value == 'IP SAN SEBASTIAN'
    print(name, 'cabecera', header_row, 'fila', row_number)
    for col in range(3, sheet.max_column + 1):
        year = sheet.cell(header_row, col).value
        if year in (2022, 2023, 2024, 2025):
            c = sheet.cell(row_number, col)
            print(year, c.coordinate, c.value)

p = ROOT / 'procesos/caracterizacion/gobernanza/sies/RETENCION/RETENCION_SIES_IPSS_162_GRANULAR_2025_CORREO_20260909.tsv'
with p.open('rb') as f:
    assert hashlib.file_digest(f, 'sha256').hexdigest() == '1e1e4872c8aa54ba97c1c4749d1fa509e13a4c11f498e933c9efddbb92b24c61'
with p.open(encoding='utf-8', newline='') as f:
    rows = list(csv.DictReader(f, delimiter='\t'))
codes = [r for r in rows if r['mc_codigo_unico'].startswith('I162')]
total = [r for r in rows if r['mc_codigo_unico'] == 'Total general']
assert len(codes) == 36 and len(total) == 1
for field in ('ret_1_retenido', 'ret_1_desertor', 'ret_1_no_considerar', 'N° de casos'):
    assert sum(int(r[field]) if r[field] not in ('-', '') else 0 for r in codes) == int(total[0][field])
print('2025 correo: 36 códigos; Total general:', {k: total[0][k] for k in ('ret_1_retenido','ret_1_desertor','ret_1_no_considerar','N° de casos')})
