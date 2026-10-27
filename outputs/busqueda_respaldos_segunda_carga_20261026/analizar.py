"""Comparar fuentes auxiliares de segunda carga sin exportar identidades."""
import collections,hashlib,json,re
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
base=Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Datos DAI (RAW) - 1.1 Datos crudos/1.2.1 Datos Institucionales/1.2.1.1 SIES/Enviados/Matrícula Unificada/2026/SEGUNDA CARGA/ARCHIVOS VARIOS')
old=Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/IPSS_CONCILIACION_UMAS_SIES_BASES/04_UMAS_PROMEDIOSDEALUMNOS/PROMEDIOSDEALUMNOS_7804.xlsx')
prior=ROOT/'outputs/busqueda_respaldo_titulos_homologaciones_20261026/RESULTADO_AGREGADO.json';cases=json.loads(Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json').read_text())
targets={(x['carrera'],x['plan_revisado'],x['ramo']) for x in json.loads(prior.read_text())['ramos']};persons={x['RUT'] for x in cases};clis={o['CODCLI'] for c in cases for o in c['contextos_v11']};assert len(targets)==21 and len(persons)==4
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
def s(x):
 if x is None:return ''
 if isinstance(x,(int,float)) and x==int(x):return str(int(x))
 return str(x).strip().upper()
def person(n,d):
 n=s(n).replace('.','').replace(' ','');d=s(d)
 return str(int(n))+'-'+d if n.isdigit() and re.fullmatch('[0-9K]',d) else None
files={'actual':(base/'promedio alumnos enviado por alvaro.xlsx','46c13bb910a9f43071583a5bc9c022be90c73188a7c29789f38f8a6f6ce3f89c'),'previo':(old,'83a0ba72b21d25dd0d151b75fb0182a1f55bbe3f5efba06d088c994067262fca')};rows={};metadata=[]
for label,(path,expected) in files.items():
 assert sha(path)==expected;w=openpyxl.load_workbook(path,read_only=True,data_only=True);it=w['Hoja1'].iter_rows(values_only=True);h={s(v):i for i,v in enumerate(next(it)) if v is not None};fields=('RUT','DIG','CODCLI','CODCARR','PLAN_DE_ESTUDIO','ANO','PERIODO','CODRAMO','CODSECC','ESTADO','DESCRIPCION_ESTADO','CONVALIDADO','RAMOEQUIV');assert set(fields).issubset(h)
 dic=collections.defaultdict(list);n=0;invalid=0
 for row in it:
  def v(k):return s(row[h[k]]) if h[k]<len(row) else ''
  career,plan,code=(v(k) for k in ('CODCARR','PLAN_DE_ESTUDIO','CODRAMO'))
  if (career,plan,code) not in targets:continue
  n+=1;id=person(v('RUT'),v('DIG'))
  if not id or not v('CODCLI') or not v('ANO') or not v('PERIODO'):invalid+=1;continue
  key=(id,v('CODCLI'),career,plan,v('ANO'),v('PERIODO'),code,v('CODSECC'));dic[key].append(tuple(v(x) for x in ('ESTADO','DESCRIPCION_ESTADO','CONVALIDADO','RAMOEQUIV')))
 w.close();assert sha(path)==expected;rows[label]=dic;metadata.append({'etiqueta':label,'ruta':str(path),'sha256':expected,'filas_objetivo':n,'llaves_distintas':len(dic),'llaves_incompletas':invalid,'multiples':sum(len(v)>1 for v in dic.values())})
a=rows['previo'];b=rows['actual'];both=set(a)&set(b);extra=set(b)-set(a);absent=set(a)-set(b);status=collections.Counter();extras=collections.Counter()
for k in extra:
 extras[(k[2],k[3],k[6])]+=1
 for x in b[k]:status[(x[0],x[1])]+=1
differences=collections.Counter()
for k in both:
 if len(a[k])==len(b[k])==1:
  for field,left,right in zip(('ESTADO','DESCRIPCION_ESTADO','CONVALIDADO','RAMOEQUIV'),a[k][0],b[k][0]):
   if left!=right:differences[field]+=1
aux=base/'enviar.xlsx';auxsha=sha(aux);w=openpyxl.load_workbook(aux,read_only=True,data_only=True);it=w['Hoja2'].iter_rows(values_only=True);head={s(v):i for i,v in enumerate(next(it)) if v is not None};assert {'N_DOC','DV','CODCLI','PLAN_DE_ESTUDIO','CODIGO_CARRERA_SIES_FINAL'}.issubset(head);matched=set();with_cli=set();attrs=collections.Counter();n=0
for row in it:
 def v(k):return s(row[head[k]]) if k in head and head[k]<len(row) else ''
 id=person(v('N_DOC'),v('DV'))
 if id not in persons:continue
 n+=1;matched.add(id)
 if v('CODCLI') in clis:with_cli.add(id);attrs['filas_codcli_coincidente']+=1
 if v('CODIGO_CARRERA_SIES_FINAL'):attrs['filas_codigo_sies_no_vacio']+=1
w.close();assert sha(aux)==auxsha
out={'clasificacion':'B fuente; C cruce; E vinculo formal y homologación general','fuentes':metadata,'comparacion':{'llaves_comunes':len(both),'solo_respaldo_previo':len(absent),'solo_respaldo_alvaro':len(extra),'ramos_con_llaves_adicionales':len(extras),'llaves_adicionales_por_ramo':[{'carrera':c,'plan':p,'ramo':r,'llaves':n} for (c,p,r),n in sorted(extras.items())],'estado_llaves_adicionales':[{'estado':a,'descripcion':b,'filas':n} for (a,b),n in sorted(status.items())],'diferencias_campos_llaves_comunes':dict(differences)},'libro_auxiliar':{'ruta':str(aux),'sha256':auxsha,'hoja':'Hoja2','personas_buscadas':4,'personas_encontradas':len(matched),'personas_con_codcli_coincidente':len(with_cli),'filas_identidad':n,'indicadores':dict(attrs),'columnas_titulo_fecha_titulo': [k for k in head if ('TITULO' in k or 'TITULACION' in k)],'limite':'Libro de matrícula; un vínculo persona-CODCLI-plan-código no prueba el plan del título ni la aceptación SIES.'},'limites':'No sumar extracciones. La fecha de extracción y completitud no están certificadas. Sin modificación de bases, reglas o denominadores.'}
(OUT/'RESULTADO_AGREGADO.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n');print(json.dumps({'fuentes':metadata,'comparacion':out['comparacion'],'auxiliar':out['libro_auxiliar']},ensure_ascii=False))
