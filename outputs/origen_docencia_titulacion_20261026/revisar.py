"""Inspección de procedencia y respaldos posteriores, sin exportar identificadores."""
from pathlib import Path
import json, hashlib, io, zipfile, re, unicodedata, signal
import xml.etree.ElementTree as ET
from datetime import datetime
import openpyxl
OUT=Path(__file__).resolve().parent
ROOT=next(Path('/Users/alexi/Library/CloudStorage').glob('OneDrive-InstitutoProfesionalSanSebasti*'))
cs=json.loads(Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json').read_text())
c=next(c for c in cs if any(x['carrera']=='TCRE' for x in c['contextos_v11']))
body,dv=c['RUT'].split('-')
cods={str(x['CODCLI']).strip().upper() for x in c['datosalumnos'] if x['CODCLI']}
def norm(v):
 if isinstance(v,float) and v.is_integer():v=int(v)
 return str(v or '').strip().upper()
def hdr(v):return ''.join(ch for ch in unicodedata.normalize('NFD',norm(v)) if not unicodedata.combining(ch))
def digest(b):return hashlib.sha256(b).hexdigest()
def timeout(*a):raise TimeoutError('lectura excedió 30 segundos')
signal.signal(signal.SIGALRM,timeout)
raw=ROOT/'Datos DAI (RAW) - 1.1 Datos crudos'/'1.2.1 Datos Institucionales'
doc=next(raw.rglob('Informe de Titulados 2025_vf.xlsx'))
paths=[doc,ROOT/'2025'/'TITULADOS_3280.xlsx',ROOT/'2025'/'Copia de TITULADOS_3280.xlsx']
for d in ROOT.iterdir():
 if d.is_dir() and (d.name.startswith('Archivos de Microsoft') or d.name=='Attachments'):
  paths.extend(p for p in d.glob('Titulados_Empleabilidad*.xlsx'))
results=[]
for p in paths:
 rec={'ruta':str(p),'clasificacion':'B','hojas':[]}
 try:
  signal.alarm(30);b=p.read_bytes();signal.alarm(0);rec['sha256']=digest(b)
  with zipfile.ZipFile(io.BytesIO(b)) as z:
   core=ET.fromstring(z.read('docProps/core.xml'))
   rec['metadatos']={n.tag.split('}')[-1]:n.text for n in core if n.tag.split('}')[-1] in ['created','modified']}
   rec['partes_procedencia']=[s for s in z.namelist() if any(t in s.lower() for t in ['connection','externallink','querytable','comment','embedding'])]
   rec['customxml_esquemas']=[]
   for n in z.namelist():
    if n.startswith('customXml/') and n.endswith('.xml'):
     el=ET.fromstring(z.read(n));rec['customxml_esquemas'].append({'parte':n,'raiz':el.tag})
  w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True)
  for s in w:
   rr=[];heads=None;headrow=None;hits=[];nf=0
   for rn,row in enumerate(s.values,1):
    if rn<=15:
     hs=[hdr(v) for v in row]
     if heads is None and any(h in ['RUT','NUM_DOCUMENTO','CODCLI','N_DOC','RUN','RUT_ALUMNO'] for h in hs):heads=hs;headrow=rn
    vals=[norm(v) for v in row]
    exactcod=any(v in cods for v in vals if v)
    full=any(re.sub(r'[.\s-]','',v)==body+dv for v in vals if v)
    bare=body in vals
    if not(exactcod or full or bare):continue
    h={'fila':rn,'codcli_exacto':exactcod,'documento_completo':full,'cuerpo_documento':bare}
    if heads:
     mapped={heads[i]:v for i,v in enumerate(row) if i<len(heads) and heads[i]}
     dvs=[norm(v) for k,v in mapped.items() if k in ['DV','DIG','DIGITO','DIGITO_VERIFICADOR']]
     h['dv_independiente_coincide']=bare and dv in dvs
     # Solamente columnas académicas; nunca nombres personales, RUT o CODCLI.
     h['campos_academicos']={k:(v.isoformat() if isinstance(v,datetime) else v) for k,v in mapped.items() if v is not None and (k in ['CODCARR','CODCARPR','CARRERA','PLAN','PLAN_DE_ESTUDIO','PLAN_ESTUDIOS','CODPESTUD','NOMBRE_TITULO','FECHA_OBT_TIT_GRA','FECHA_TITULACION','FECHA','ESTADO_ACADEMICO','SITUACION_ACADEMICA'] or k.startswith('ANIO_INGRESO'))}
    hits.append(h)
   rec['hojas'].append({'hoja':s.title,'filas':s.max_row,'columnas':s.max_column,'fila_encabezado':headrow,'encabezados':heads,'coincidencias':hits})
  w.close();rec['integridad_conservada']=digest(p.read_bytes())==rec['sha256']
 except Exception as e:rec['error']=type(e).__name__+': '+str(e)
 signal.alarm(0);results.append(rec)
 print(p.name,[(s['hoja'],len(s['coincidencias'])) for s in rec['hojas']],rec.get('error',''),flush=True)
payload=json.dumps(results,ensure_ascii=False,indent=2,default=str)+'\n'
for x in cs:
 assert x['RUT'].split('-')[0] not in payload
 for a in x['datosalumnos']:
  if a['CODCLI']:assert norm(a['CODCLI']) not in payload
(OUT/'RESULTADOS.json').write_text(payload)
