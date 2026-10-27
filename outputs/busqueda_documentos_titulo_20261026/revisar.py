"""Búsqueda acotada de contenido, solo lectura; sin exportar identidades."""
import collections, hashlib, io, json, re, signal, unicodedata, zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from pypdf import PdfReader
import openpyxl

OUT = Path(__file__).resolve().parent
PRIVATE = Path('/Users/alexi/.local/share/ipss/busqueda_documentos_titulo_20261026')
PRIVATE.mkdir(mode=0o700, parents=True, exist_ok=True)
ROOT = next(Path('/Users/alexi/Library/CloudStorage').glob('OneDrive-InstitutoProfesionalSanSebasti*'))
CASES = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
cases = json.loads(CASES.read_text())
case = next(c for c in cases if any(x['carrera']=='TCRE' for x in c['contextos_v11']))
body, dv = case['RUT'].split('-')
codclis = {str(x['CODCLI']).strip().upper() for x in case['datosalumnos'] if x['CODCLI']}

def norm(value):
    if isinstance(value,float) and value.is_integer(): value=int(value)
    return ''.join(c for c in unicodedata.normalize('NFD',str(value or '').upper()) if not unicodedata.combining(c))

def find(text):
    text=norm(text)
    full = re.search(r'(?<![0-9])'+r'[.\s]*'.join(body)+r'\s*-?\s*'+dv+r'(?![A-Z0-9])',text) is not None
    return {'codcli':any(re.search(r'(?<![A-Z0-9])'+re.escape(c)+r'(?![A-Z0-9])',text) for c in codclis),
            'rut_dv':full,'cuerpo':re.search(r'(?<![0-9])'+r'[.\s]*'.join(body)+r'(?![0-9])',text) is not None}

def timeout(*_): raise TimeoutError('lectura superior a 8 segundos')
signal.signal(signal.SIGALRM, timeout)
roots=[ROOT/'0 BASES DE DATOS',ROOT/'Datos DAI (RAW) - 1.1 Datos crudos'/'1.2.1 Datos Institucionales',ROOT/'Archivos de chat de Microsoft Teams']
paths=sorted({p for r in roots for p in r.rglob('*') if p.is_file() and p.suffix.lower() in ('.pdf','.docx') and not p.name.startswith('~$')})
teams=ROOT/'Archivos de chat de Microsoft Teams'
paths += [teams/n for n in ['TITULADOS_3280.xlsx','TITULADOS_3280 (1).xlsx','TITULADOS_3280 1.xlsx','Ingenieria ciber_redes_titulados_06agosto.xlsx'] if (teams/n).exists()]
results=[]; seen={}
for index,p in enumerate(paths):
    rec={'ruta':str(p),'formato':p.suffix.lower()}
    try:
        signal.alarm(8)
        blob=p.read_bytes(); digest=hashlib.sha256(blob).hexdigest(); rec['sha256']=digest
        if digest in seen:
            rec.update(estado='duplicado_por_hash',igual_a=seen[digest]); results.append(rec); continue
        seen[digest]=str(p)
        hits=[]; empty=[]; units=0
        if p.suffix.lower()=='.xlsx':
            book=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
            rec['hojas']=book.sheetnames; rec['cabeceras_relevantes']={}
            for ws in book:
                headers=[]
                for rownum,row in enumerate(ws.iter_rows(values_only=True),1):
                    units+=1
                    if rownum<=15:
                        for v in row:
                            nv=norm(v)
                            if nv in ('CODCLI','CODCARPR','PLAN_DE_ESTUDIO','PLAN_ESTUDIOS','CODPESTUD','FECHA','FECHA_TITULO','TITULO','NUMERO_REGISTRO','ANO','PERIODO','CARRERA'):
                                headers.append(nv)
                    flags=find(' | '.join(norm(v) for v in row))
                    if any(flags.values()):hits.append({'hoja':ws.title,'fila':rownum,**flags})
                rec['cabeceras_relevantes'][ws.title]=sorted(set(headers))
            book.close()
        else:
            if p.suffix.lower()=='.pdf':
                pdf=PdfReader(io.BytesIO(blob));texts=[page.extract_text() or '' for page in pdf.pages]
            else:
                with zipfile.ZipFile(io.BytesIO(blob)) as z:
                    texts=[' '.join(ET.fromstring(z.read(n)).itertext()) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml') and any(k in n for k in ['document.xml','header','footer','comments','footnotes','endnotes'])]
            units=len(texts)
            for i,t in enumerate(texts,1):
                if not t.strip():empty.append(i)
                flags=find(t)
                if any(flags.values()):hits.append({'unidad':i,**flags})
        rec.update(estado='leido',unidades=units,unidades_sin_texto=empty,coincidencias=hits)
        rec['hash_conservado']=hashlib.sha256(p.read_bytes()).hexdigest()==digest
    except Exception as exc:
        rec.update(estado='pendiente_lectura',error=type(exc).__name__)
    finally:
        signal.alarm(0)
        if not results or results[-1] is not rec:results.append(rec)
    if (index+1)%10==0:print(json.dumps({'revisados':index+1,'inventariados':len(paths)}),flush=True)

payload=json.dumps(results,ensure_ascii=False,indent=2)+'\n'
for c in cases:
    assert c['RUT'].split('-')[0] not in payload
    for x in c['datosalumnos']:
        if x['CODCLI']:assert str(x['CODCLI']) not in payload
dest=PRIVATE/'INVENTARIO.json';dest.write_text(payload);dest.chmod(0o600)
summary={'clasificacion':'B fuentes; C búsqueda; E atribución pendiente','fecha':datetime.now().astimezone().isoformat(),
         'fuentes':len(results),'formatos':dict(collections.Counter(x['formato'] for x in results)),
         'estados':dict(collections.Counter(x['estado'] for x in results)),
         'fuentes_con_coincidencias':sum(bool(x.get('coincidencias')) for x in results),
         'fuentes_con_unidades_sin_texto':sum(bool(x.get('unidades_sin_texto')) for x in results),
         'hashes_verificados_sin_cambio':sum(x.get('hash_conservado',False) for x in results),
         'limite':'PDF sin OCR; DOCX XML textual; Excel valores guardados, no recálculo. Coincidencias son pistas, no atribución acreditada. Duplicados se leen una vez. Ocho segundos por archivo.',
         'inventario_privado':str(dest)}
(OUT/'RESULTADO_AGREGADO.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(summary,ensure_ascii=False),flush=True)
for x in results:
    if x.get('coincidencias'):print(json.dumps({'formato':x['formato'],'sha256':x['sha256'],'unidades_coincidentes':len(x['coincidencias'])}))
