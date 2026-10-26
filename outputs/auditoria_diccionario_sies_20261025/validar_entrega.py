"""Valida el Excel exportado, privacidad, integridad y copia final al Escritorio."""
import collections
import hashlib
import json
import math
import re
import shutil
import sqlite3
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
FILE = OUT/'DICCIONARIO_UMAS_SIES_RESUMEN_VERIFICADO_20261025.xlsx'
DEST = Path.home()/'Desktop'/FILE.name
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
q = json.loads((OUT/'QA_CONSULTA.json').read_text())
source = json.loads((OUT/'insumos/resumen.json').read_text())
desc = json.loads(Path(source['descriptor']).read_text())
wb = openpyxl.load_workbook(FILE,data_only=True)
wf = openpyxl.load_workbook(FILE,data_only=False)
checks = []
def check(name, result):
    checks.append({'control':name,'correcto':bool(result),'clase':'C'})
    if not result:
        raise AssertionError(name)

expected={'C8':2854,'D8':2737,'E8':13,'F8':2750,'G8':104,'H8':2750/2854,
          'C9':1077,'F9':1053,'G9':24,'H9':1053/1077,
          'C10':24774,'D10':21561,'E10':164,'F10':21725,'G10':3049,'H10':21725/24774,
          'C17':26,'D17':10,'E17':16,'F17':0,'G17':16,
          'C18':93,'D18':30,'E18':59,'F18':4,'G18':63,'C30':104}
for cell,value in expected.items():
    got=wb['Resumen'][cell].value
    check('Resumen '+cell, isinstance(got,(int,float)) and math.isclose(got,value,rel_tol=1e-12))
for cell in ['C29','C30','C31','C32']:
    check('Conciliación '+cell, wb['Gobierno y QA'][cell].value==0)
formula_count=0
errors=[]
for ws in wf:
    cached=wb[ws.title]
    for row in ws:
        for cell in row:
            if cell.data_type=='f':
                formula_count+=1
                if cached[cell.coordinate].value is None or cached[cell.coordinate].data_type=='e':
                    errors.append(ws.title+'!'+cell.coordinate)
check('Fórmulas con caché sin errores',not errors)
for ws in wf:
    for table in ws.tables.values():
        mincol,minrow,maxcol,maxrow=openpyxl.utils.range_boundaries(table.ref)
        headers=[ws.cell(minrow,c).value for c in range(mincol,maxcol+1)]
        check('Encabezados tabla '+table.name,headers==[c.name for c in table.tableColumns] and len(set(headers))==len(headers))
        check('Filtro tabla '+table.name,table.autoFilter is not None)
with zipfile.ZipFile(FILE) as z:
    check('ZIP íntegro',z.testzip() is None)
    xml={n:z.read(n) for n in z.namelist() if n.endswith('.xml')}
    for n,content in xml.items():ET.fromstring(content)
    check('XML válido',True)
    check('Sin vínculos externos',not any(n.startswith('xl/externalLinks/') for n in z.namelist()))
    db=sqlite3.connect(Path(desc['base_privada']['ruta']).as_uri()+'?mode=ro',uri=True)
    secrets=set()
    for persona,codcli in db.execute('SELECT DISTINCT persona,codcli FROM observaciones'):
        secrets.update([persona,codcli])
    db.close()
    package='\n'.join(b.decode('utf8') for b in xml.values())
    check('Sin identificadores personales del diccionario',not any(s and s in package for s in secrets))
check('Base privada conserva hash',sha(desc['base_privada']['ruta'])==desc['base_privada']['sha256'])
check('V17 conserva hash',sha(Path.home()/'Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261024_V17.xlsx')==q['v17_escritorio_sha256'])
head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
check('HEAD conservado',head==q['head'])
check('Cambios pendientes conservados',all(sha(ROOT/p)==h for p,h in q['cambios_previos_sha256'].items()))
if DEST.exists():
    check('Copia previa compatible',sha(DEST)==sha(FILE))
else:
    shutil.copy2(FILE,DEST)
check('Copia Escritorio idéntica',sha(FILE)==sha(DEST))
result={'estado':'QA_CORRECTO_CON_LIMITES_EXPLICITOS','clasificacion':'B/C verificación; E límites',
        'controles':checks,'formulas':formula_count,'errores_formula':errors,
        'hojas':wb.sheetnames,'tablas':sum(len(ws.tables) for ws in wf),
        'archivo':str(FILE),'escritorio':str(DEST),'sha256':sha(FILE),
        'rama':subprocess.check_output(['git','branch','--show-current'],cwd=ROOT,text=True).strip(),
        'HEAD_antes':q['head'],'HEAD_despues':head,
        'limites':['Recálculo Artifact Tool y lectura independiente OpenXML/openpyxl; no se abrió en Excel nativo.',
                   'La paridad del servicio acredita reproducción, no verdad documental universal ni aceptación SIES.',
                   'No se cerraron pendientes ni se publicaron cambios en GitHub, complementos o skills.'],
        'commit_push':False,'bases_y_denominadores_modificados':False}
(OUT/'QA_ENTREGA.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
inputs={str(p.relative_to(OUT)):sha(p) for p in sorted((OUT/'insumos').glob('*.json'))}
manifest={'recurso':'Auditoría ejecutiva UMAS a código único SIES','fecha':'2026-10-25',
          'fuente_activa':source['descriptor'],'edicion':source['edicion'],'inputs_agregados_sha256':inputs,
          'archivos_salida':{str(FILE):sha(FILE),str(DEST):sha(DEST)},
          'codigos_documentados':21561,'correspondencias_internas':164,'pendientes':3049,
          'combinaciones_historicas':{'definicion':'carrera+plan+sede+jornada','total':93,'completas':30,'parciales':59,'sin_resolver':4},
          'v17':{'contextos':2854,'con_codigo':2750,'pendientes':104,'combinaciones':26,'combinaciones_con_pendiente':16},
          'clasificacion':'B fuentes y estados; C cálculos/QA; D decisiones conservadas; E pendientes. A solo paleta de marca.',
          'politica':'Solo agregados, sin identificadores. Fuentes individuales permanecen fuera de Git. No modifica el diccionario ni los denominadores.',
          'estado_git':{'rama':result['rama'],'head':head,'publicado':False,'commit_push':False},
          'reproducibilidad':'auditar_cobertura.py y auditar_capas.py leen fuentes locales en modo lectura; verificar_consulta.py valida consumidores y hashes; build.mjs crea el reporte desde agregados; validar_entrega.py verifica y copia. Se necesita acceso local a las fuentes privadas; no se empaquetan.',
          'qa':str(OUT/'QA_ENTREGA.json')}
(OUT/'MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({k:v for k,v in result.items() if k not in ('controles','limites')},ensure_ascii=False))
