"""Inyecta solo celdas agregadas y caches recalculadas; conserva las demás partes."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
import json,copy,hashlib,re,shutil
import openpyxl
O=Path(__file__).parent;P=Path('/Users/alexi/.local/share/ipss/fase3_candidata_20261025');src=Path('/Users/alexi/Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261024_V17.xlsx');dst=P/'RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261025_V18_CANDIDATA.xlsx'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest();assert sha(src)=='8b461d53f75ac26e2e38aa46d0aae921736d21ac2a64034cfd26c08495756a52';assert not dst.exists()
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda t:'{'+N+'}'+t
with ZipFile(src) as z:base={n:z.read(n) for n in z.namelist()}
with ZipFile(P/'donante.xlsx') as z:donor={n:z.read(n) for n in z.namelist()}
strings=E.fromstring(donor['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in donor else []
def converted(cell):
 c=copy.deepcopy(cell)
 if c.get('t')=='s':
  si=strings[int(c.find(q('v')).text)];c.remove(c.find(q('v')));c.set('t','inlineStr');el=E.SubElement(c,q('is'))
  for ch in si:el.append(copy.deepcopy(ch))
 return c
def replace_content(old,new):
 for ch in list(old):old.remove(ch)
 if 't' in new.attrib:old.set('t',new.get('t'))
 elif 't' in old.attrib:del old.attrib['t']
 for ch in new:old.append(copy.deepcopy(ch))
s52=E.fromstring(base['xl/worksheets/sheet52.xml']);s54=E.fromstring(base['xl/worksheets/sheet54.xml'])
d52={c.get('r'):converted(c) for c in E.fromstring(donor['xl/worksheets/sheet1.xml']).iter(q('c'))};d54={c.get('r'):converted(c) for c in E.fromstring(donor['xl/worksheets/sheet3.xml']).iter(q('c'))}
notes=json.loads((O/'NOTAS_EXCEL.json').read_text());changed52=[]
for c in s52.iter(q('c')):
 r=c.get('r')
 if r in notes:replace_content(c,d52[r]);changed52.append(r)
 elif c.find(q('f')) is not None and r in d52:
  # Transfer cached value only, never rewrite original formula.
  old=c.find(q('v'));new=d52[r].find(q('v'))
  if new is not None and (old is None or old.text!=new.text):
   assert re.fullmatch(r'[FHI](17|18|19|20|21|22|24)',r),r
   if old is not None:c.remove(old)
   c.append(copy.deepcopy(new));changed52.append(r)
# Use original row styles, selecting documentary/pending style by corresponding old status.
w=openpyxl.load_workbook(src,read_only=True,data_only=True);rows54=list(w.worksheets[53].iter_rows(min_row=6,max_row=228,max_col=15,values_only=True));w.close();styleby={}
for row in s54.find(q('sheetData')):
 rn=int(row.get('r'))
 if rn>=6 and rn<=228:
  st=rows54[rn-6][8];styleby.setdefault(st,{re.sub(r'\d','',c.get('r')):c.get('s') for c in row})
newtable=json.loads((O/'ENTRADAS_EXCEL.json').read_text())['tabla54_candidata']
for c in s54.iter(q('c')):
 r=c.get('r');rn=int(re.search(r'\d+',r).group());col=re.sub(r'\d','',r)
 if r=='A1':replace_content(c,d54[r])
 elif 6<=rn<=228:
  if r in d54:replace_content(c,d54[r])
  else:
   for ch in list(c):c.remove(ch)
   c.attrib.pop('t',None)
  if rn-6<len(newtable):
   st=newtable[rn-6][8];style=styleby[st].get(col)
   if style is not None:c.set('s',style)
enc=lambda x:E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)
parts={'xl/worksheets/sheet52.xml':enc(s52),'xl/worksheets/sheet54.xml':enc(s54)}
with ZipFile(dst,'w',ZIP_DEFLATED) as z:
 for n,v in base.items():z.writestr(n,parts.get(n,v))
with ZipFile(dst) as z:
 assert z.testzip() is None
 assert set(z.namelist())==set(base)
 changedparts=[n for n in base if z.read(n)!=base[n]];assert set(changedparts)==set(parts)
 for n in z.namelist():
  if n.endswith(('.xml','.rels')):E.fromstring(z.read(n))
f=openpyxl.load_workbook(dst,read_only=True,data_only=False);v=openpyxl.load_workbook(dst,read_only=True,data_only=True);of=openpyxl.load_workbook(src,read_only=True,data_only=False);ov=openpyxl.load_workbook(src,read_only=True,data_only=True)
assert f.sheetnames==of.sheetnames and len(f.sheetnames)==55
# Formula text preserved everywhere, unrelated sheets byte-identical.
for si in [51,53]:
 before={c.coordinate:c.value for row in of.worksheets[si] for c in row if c.data_type=='f'};after={c.coordinate:c.value for row in f.worksheets[si] for c in row if c.data_type=='f'};assert before==after
assert [list(row) for row in v.worksheets[53].iter_rows(min_row=6,max_row=226,max_col=15,values_only=True)]==newtable
assert all(c.value is None for row in v.worksheets[53].iter_rows(min_row=227,max_row=228,max_col=15) for c in row)
s=v.worksheets[51];checks=0
for rn in range(17,23):
 year=s[f'A{rn}'].value;period=s[f'C{rn}'].value
 for col,state in [('F','B/C Documental anual'),('G','D Interno aprobado'),('H','E Pendiente')]:
  want=sum(x[13] for x in newtable if x[0]==year and x[2]==period and x[8]==state);assert s[f'{col}{rn}'].value==want;checks+=1
 assert s[f'I{rn}'].value==sum(s[f'{c}{rn}'].value for c in 'FGH');checks+=1
assert [s[f'{c}24'].value for c in 'FGHI']==[2742,13,99,2854]
for row in range(6,12):
 for col in 'DEFG':assert s[f'{col}{row}'].value==ov.worksheets[51][f'{col}{row}'].value
errors=[(sheet.title,c.coordinate,c.value) for sheet in [v.worksheets[51],v.worksheets[53]] for row in sheet for c in row if c.data_type=='e'];assert not errors
for book in [f,v,of,ov]:book.close()
assert sha(src)=='8b461d53f75ac26e2e38aa46d0aae921736d21ac2a64034cfd26c08495756a52'
desktop=Path('/Users/alexi/Desktop')/dst.name;assert not desktop.exists();shutil.copy2(dst,desktop);assert sha(dst)==sha(desktop)
qa={'clasificacion':'C','estado':'QA_ESTRUCTURAL_Y_RECALCULO_AFECTADO_CORRECTO','archivo':str(dst),'escritorio':str(desktop),'sha256':sha(dst),'hojas':55,'hojas_bytes_conservados':53,'graficos_bytes_conservados':sum(bool(re.fullmatch(r'xl/charts/chart\d+.xml',n)) for n in base),'tablas_bytes_conservadas':sum(bool(re.fullmatch(r'xl/tables/table\d+.xml',n)) for n in base),'estilos_originales_conservados':True,'formulas_originales_conservadas':True,'tasas_y_denominadores_conservados':True,'controles_aritmeticos':checks+4,'partes_cambiadas':changedparts,'celdas52_actualizadas':changed52,'errores_hojas_afectadas':errors,'recálculo':'Artifact Tool, dependencias agregadas afectadas; demás caches conservados','limite':'No prueba apertura ni recálculo completo en Excel nativo. La copia hereda las hojas individuales ya presentes: no publicar el XLSX en Git.'}
(O/'QA_EXCEL.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa,ensure_ascii=False))
