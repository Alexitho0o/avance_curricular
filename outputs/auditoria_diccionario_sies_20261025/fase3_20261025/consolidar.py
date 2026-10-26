"""Candidata inmutable: ocho reparaciones exactas, QA global y agregados Excel."""
import json,sqlite3,hashlib,copy,collections,os,subprocess,sys,datetime
from pathlib import Path
import openpyxl
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');O=Path(__file__).parent;P=Path('/Users/alexi/.local/share/ipss/fase3_candidata_20261025');F2=P.parent/'fase2_historico_20261025'
sys.path.insert(0,str(R));from common.diccionario_sies_seguimiento import DiccionarioSiesSeguimiento
read=lambda p:json.loads(Path(p).read_text());sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
assert subprocess.run(['git','-C',str(P),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
active=R/'data/diccionario_sies_seguimiento/activo.json';a=read(active);d=read(R/a['descriptor']);assert sha(R/a['descriptor'])==a['sha256']
assert sha(d['base_privada']['ruta'])==d['base_privada']['sha256']
ops=read(F2/'OPERACIONES_CANDIDATAS.json')['operaciones'];assert len(ops)==len({x['observacion_id'] for x in ops})==8
allowed=set('codigos_anuales_documentados codigo_documental_atribuible codigo_interno_adoptado estado_documental estado_interno regla clasificacion requisito_pendiente evidencias_matricula'.split())
path=P/'DICCIONARIO_SIES_CANDIDATO.sqlite';assert not path.exists()
old=sqlite3.connect('file:'+d['base_privada']['ruta']+'?mode=ro',uri=True);new=sqlite3.connect(path);old.backup(new)
mu_before=new.execute('select count(*) from mu_anual').fetchone()[0]
for op in ops:
 oid=op['observacion_id'];before=op['estado_antes'];after=op['estado_despues_propuesto'];current=json.loads(new.execute('select payload from observaciones where id=?',(oid,)).fetchone()[0]);assert current==before
 assert set(before)==set(after) and {k for k in before if before[k]!=after[k]}<=allowed
 assert after['codigo_documental_atribuible']==after['codigo_interno_adoptado'] and not before['codigo_documental_atribuible'] and not before['codigo_interno_adoptado']
 new.execute('update observaciones set payload=? where id=?',(json.dumps(after,ensure_ascii=False),oid))
 for e in after['evidencias_matricula']:
  payload={**e,'codigo':e['codigo_SIES_observado'],'anio':e['ANIO_INFORMADO'],'persona':after['persona']}
  new.execute('insert or ignore into mu_anual values (?,?,?,?,?,?,?)',(after['persona'],payload['anio'],payload['codigo'],e['fuente'],e['hoja'],e['fila'],json.dumps(payload,ensure_ascii=False)))
new.commit();assert new.execute('pragma integrity_check').fetchone()[0]=='ok'
oldrows={r[0]:r for r in old.execute('select * from observaciones')};newrows={r[0]:r for r in new.execute('select * from observaciones')};changed={k for k in oldrows if oldrows[k]!=newrows[k]};assert changed=={x['observacion_id'] for x in ops}
assert len(oldrows)==len(newrows)==24774
oldmu=list(old.execute('select * from mu_anual'));assert all(new.execute('select * from mu_anual where persona=? and anio=? and codigo=? and fuente=? and hoja=? and fila=?',r[:6]).fetchone()==r for r in oldmu)
mu_added=new.execute('select count(*) from mu_anual').fetchone()[0]-mu_before
rows=[json.loads(r[-1]) for r in newrows.values()];counts=collections.Counter('documental' if x['codigo_documental_atribuible'] else 'interno' if x['codigo_interno_adoptado'] else 'pendiente' for x in rows)
assert counts=={'documental':21569,'interno':164,'pendiente':3041}
new.close();old.close();os.chmod(path,0o600)
nd=copy.deepcopy(d);nd['edicion']='v8_candidata_20261025_ocho_p';nd['estado']='CANDIDATA_NO_ACTIVA';nd['fecha_hora']=datetime.datetime.now().astimezone().isoformat();nd['base_privada']={'ruta':str(path),'sha256':sha(path)};nd['edicion_anterior']={'descriptor':a['descriptor'],'sha256':a['sha256']};nd.pop('adjudicaciones_privadas',None)
nd['reparaciones_C']={'ruta':str(F2/'OPERACIONES_CANDIDATAS.json'),'sha256':sha(F2/'OPERACIONES_CANDIDATAS.json'),'operaciones':8,'regla_existente':True};nd['resumen']['codigo_documental_reutilizado']=21569;nd['resumen']['pendientes_internos']={'observaciones':3041,'personas':len({x['persona'] for x in rows if not(x['codigo_documental_atribuible'] or x['codigo_interno_adoptado'])})};nd['resumen']['pendientes_documentales']={'observaciones':3205,'personas':len({x['persona'] for x in rows if not x['codigo_documental_atribuible']})};nd['resumen']['nuevas_asignaciones']=8
# Recompute stale candidate counts rather than carry v7 scope metrics as current.
for key in ['filas_indice_mu_total_historico','filas_indice_mu_2022_2026','puentes_documentales_explicitos','filas_mu_sin_llave_rut_en_indice','foco_versiones_multiples','excepciones_por_estado']:nd['resumen'].pop(key,None)
nd['resumen']['filas_indice_mu_total_historico']=mu_before+mu_added;nd['resumen']['enlaces_mu_anadidos']=mu_added
nd['resumen']['excepciones_por_estado']=dict(collections.Counter(x['estado_interno'] for x in rows if not(x['codigo_documental_atribuible'] or x['codigo_interno_adoptado'])))
write(O/'descriptor_candidato.json',nd)
service=DiccionarioSiesSeguimiento(descriptor=O/'descriptor_candidato.json');base=DiccionarioSiesSeguimiento();querydelta=[]
for x in rows:
 q=dict(persona=x['persona'],codcli=x['CODCLI'],plan=x['plan'],anio=x['ano'],periodo=x['periodo'],carrera=x['carrera'],sede=x['sede_dimension'],jornada=x['jornada_sies'])
 n=service.consultar(**q);b=base.consultar(**q);assert n['codigo_sies']==(x['codigo_documental_atribuible'] or x['codigo_interno_adoptado'] or None)
 n.pop('edicion');b.pop('edicion')
 if n!=b:querydelta.append(x['observacion_id'])
assert set(querydelta)==changed
service.close();base.close()
# Reproduce old aggregate exactly before replacing only five context attributions.
ctxpath=P.parent/'destinos_calendario_v7_20261023/CONTEXTOS_POR_ORIGEN_PERIODO.json';ctx=read(ctxpath);ori=read(P.parent/'aplicacion_diccionario_intertrimestral_20261023_125556/ORIGENES_T1.json');byrow={x['fila_foto']:x for x in ori};index={x['observacion_id']:x for x in rows}
def aggregate(contexts):
 cnt=collections.Counter();states={'B/C_DOCUMENTAL_ANUAL_ATRIBUIDO':'B/C Documental anual','D_CORRESPONDENCIA_INTERNA_APROBADA':'D Interno aprobado','E_PENDIENTE':'E Pendiente'}
 for x in contexts:
  o=byrow[x['fila_origen_foto']];s=states[x['estado_destino']];lim='Atribución documental al período pendiente' if s.startswith('D') else 'Versión/atribución pendiente; no adjudicar por totales' if s.startswith('E') else 'Código anual atribuido; no acredita día del cambio ni aceptación SIES'
  cnt[(x['anio_origen'],x['anio_observado'],int(x['periodo_nativo']),o['carrera'],x['carrera_umas'],x['plan'],x['codigo_origen_T1'],x['codigo_destino'] or 'Pendiente',s,x['tipo_pendiente'] or x['relacion_codigo'],x['comparacion_plan'],x['contraste_dimensiones'],'Sí' if o['multiple'] else 'No',lim)]+=1
 return [list(k[:13])+[v,k[13]] for k,v in sorted(cnt.items(),key=lambda z:tuple(str(a) for a in z[0]))]
wbpath=Path('/Users/alexi/Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261024_V17.xlsx');assert sha(wbpath)=='8b461d53f75ac26e2e38aa46d0aae921736d21ac2a64034cfd26c08495756a52'
w=openpyxl.load_workbook(wbpath,read_only=True,data_only=False);oldtable=[list(r) for r in w.worksheets[53].iter_rows(min_row=6,max_row=228,max_col=15,values_only=True)];assert oldtable==aggregate(ctx)
newctx=copy.deepcopy(ctx);nchanges=0
for x in newctx:
 if x['observacion_id'] in changed:
  r=index[x['observacion_id']];x.update(codigo_destino=r['codigo_documental_atribuible'],codigo_documental=r['codigo_documental_atribuible'],codigo_interno=r['codigo_interno_adoptado'],estado_destino='B/C_DOCUMENTAL_ANUAL_ATRIBUIDO',regla=r['regla'],pendiente=r['requisito_pendiente'],tipo_pendiente=None,relacion_codigo='MISMO_CODIGO_T1' if x['codigo_origen_T1']==r['codigo_documental_atribuible'] else 'CODIGO_DISTINTO_RESPALDADO');nchanges+=1
assert nchanges==5
# Use exact established spelling for relation; no new categories.
relations=collections.Counter(x['relacion_codigo'] for x in ctx if x['codigo_destino'] and x['codigo_destino']==x['codigo_origen_T1']);assert len(relations)==1
same=next(iter(relations))
diffrel={x['relacion_codigo'] for x in ctx if x['codigo_destino'] and x['codigo_destino']!=x['codigo_origen_T1']};assert diffrel
for x in newctx:
 if x['observacion_id'] in changed:x['relacion_codigo']=same if x['codigo_destino']==x['codigo_origen_T1'] else (_ for _ in ()).throw(ValueError('Cambio de familia requiere clasificacion explicita'))
table=aggregate(newctx);assert sum(r[13] for r in table)==2854
privateids={x['persona'] for x in rows}|{x['CODCLI'] for x in rows}
assert all(str(v) not in privateids for row in table for v in row)
write(P/'CONTEXTOS_CANDIDATOS.json',newctx);os.chmod(P/'CONTEXTOS_CANDIDATOS.json',0o600)
# All strings and formulas from only the three aggregate sheets, no personal data.
sheets={s.title:[list(r) for r in s.iter_rows(values_only=True)] for s in w.worksheets[51:54]};w.close()
write(O/'ENTRADAS_EXCEL.json',{'hojas':sheets,'tabla54_candidata':table,'conteos_contextos':dict(collections.Counter(x['estado_destino'] for x in newctx)),'relaciones':dict(collections.Counter(x['relacion_codigo'] for x in newctx)),'multiversion':dict(collections.Counter(x['estado_destino'] for x in newctx if byrow[x['fila_origen_foto']]['multiple']))})
write(O/'QA_DICCIONARIO.json',{'clasificacion':'B/C','observaciones':24774,'consultas_verificadas':24774,'cambios_esperados':8,'sin_cambios':24766,'cambios_inesperados':0,'indice_mu_enlaces_nuevos':mu_added,'indice_mu_previo_preservado':True,'resumen':dict(counts),'contextos_v17_cambiados':5,'puntero_activo_preservado':sha(active)== 'a041fa4d8c1cd25006959fc2b05089eab396f9b684cb35c942e7263343890497','excel_original_sha256':sha(wbpath),'candidate_only':True})
print(json.dumps({'historico':dict(counts),'mu_enlaces_nuevos':mu_added,'v17_cambios':nchanges,'filas54_antes':len(oldtable),'filas54_despues':len(table)}))
