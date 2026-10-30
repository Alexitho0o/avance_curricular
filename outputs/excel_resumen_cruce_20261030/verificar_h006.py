"""QA H006: preservación íntegra, cifras nuevas, filtros y formato centrado.

openpyxl solo lee; los ensayos de selectores editan copias OOXML privadas.
Las cachés LibreOffice se inyectan sin sustituir tablas ni gráficos originales.
"""
from pathlib import Path
import json,math,re,shutil,subprocess,zipfile
import xml.etree.ElementTree as ET
import openpyxl,pypdfium2
from construir import OUT,ROOT,RUNTIME,BOOK,sha,dump,fuentes
from verificar import lo,rewrite,setup_panels,tag

PRIVATE=Path.home()/'.local/share/ipss/h006_20261030'
def same(a,b):
    if a in [None,''] and b in [None,'']:return True
    if isinstance(a,(float,int)) and isinstance(b,(float,int)):return math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-10)
    return a==b
def inject(book,recalc):
    repl={}
    with zipfile.ZipFile(book) as a,zipfile.ZipFile(recalc) as b:
        for n in a.namelist():
            if not re.fullmatch(r'xl/worksheets/sheet\d+.xml',n):continue
            r=ET.fromstring(a.read(n));br=ET.fromstring(b.read(n));bc={c.get('r'):c for c in br.iter(tag('c'))}
            for c in r.iter(tag('c')):
                if c.find(tag('f')) is None:continue
                other=bc[c.get('r')]
                for old in [tag('v'),tag('is')]:
                    e=c.find(old)
                    if e is not None:c.remove(e)
                typ=other.get('t');val=other.find(tag('v'))
                if typ=='inlineStr':
                    val=ET.Element(tag('v'));val.text=''.join(other.itertext());typ='str'
                if val is None:val=ET.Element(tag('v'));val.text='';typ='str'
                if typ:c.set('t',typ)
                else:c.attrib.pop('t',None)
                c.append(ET.fromstring(ET.tostring(val)))
            repl[n]=ET.tostring(r,encoding='utf-8',xml_declaration=True)
        wr=ET.fromstring(a.read('xl/workbook.xml'));calc=wr.find(tag('calcPr'))
        if calc is None:calc=ET.SubElement(wr,tag('calcPr'))
        calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1');calc.set('calcMode','auto');repl['xl/workbook.xml']=ET.tostring(wr,encoding='utf-8',xml_declaration=True)
    rewrite(book,repl)

def workbook_checks(book,baseline,controls,work):
    recalc,_=lo(book,work,'xlsx');fw=openpyxl.load_workbook(book,data_only=False);vw=openpyxl.load_workbook(recalc,data_only=True)
    result=[];diff=[]
    for x in controls:
        value=vw[x['hoja']][x['celda']].value;ok=same(value,x['esperado']);r=dict(x,obtenido=value,coincide=ok);result.append(r)
        if not ok:diff.append(r)
    dump(OUT/(book.stem+'_CONTROL_H006.json'),result)
    if diff:dump(OUT/'DIFERENCIAS_H006.json',diff);raise ValueError('Cifras: '+str(diff[:3]))
    errors=[];notcenter=[];formulas=0;populated=0
    for s in fw:
        assert s.sheet_state=='visible'
        for row in s:
            for c in row:
                if c.value is None:continue
                populated+=1
                if c.data_type=='f':formulas+=1
                if vw[s.title][c.coordinate].data_type=='e':errors.append([s.title,c.coordinate])
                if c.alignment.horizontal!='center' or c.alignment.vertical!='center':notcenter.append([s.title,c.coordinate,str(c.alignment)])
        for t in s.tables.values():assert t.autoFilter is not None
    assert not errors,errors
    assert not notcenter,notcenter[:3]
    oldf=openpyxl.load_workbook(baseline,data_only=False);oldv=openpyxl.load_workbook(baseline,data_only=True);oldchecked=0;changed=[]
    for s in oldf:
        for row in s:
            for c in row:
                if c.value is None:continue
                n=fw[s.title][c.coordinate];v=vw[s.title][c.coordinate]
                if c.value!=n.value or not same(oldv[s.title][c.coordinate].value,v.value):changed.append([s.title,c.coordinate,c.value,n.value])
                oldchecked+=1
    assert not changed,changed[:3]
    inject(book,recalc)
    return dict(controles=len(controls),diferencias=0,formulas=formulas,errores_formula=errors,celdas_previas_sin_cambios=oldchecked,celdas_contenido_centradas=populated,cambios_permitidos='Centrado horizontal/vertical, ajuste de texto; celdas y hojas nuevas H006 solo en resumen',hojas=fw.sheetnames,tablas=sum(len(s.tables) for s in fw),graficos=sum(len(s._charts) for s in fw)),fw

def filters(book,model,data):
    specs={s['name']:s for s in model['sheets'] if s.get('selectors')};results=[]
    # Cada dimensión varía por separado y también en combinación.
    for trial in range(4):
        choices={}
        for name,s in specs.items():
            values=['Todas']*3
            for i,sel in enumerate(s['selectors']):
                if trial==i or trial==3:
                    vs=sel['values'][1:];values[i]=vs[-1] if vs else 'Todas'
            choices[name]=values
        p=PRIVATE/f'prueba_filtros_{trial}.xlsx';shutil.copy2(book,p);repl={}
        with zipfile.ZipFile(p) as z:
            for i,s in enumerate(model['sheets'],1):
                if s['name'] not in choices:continue
                n=f'xl/worksheets/sheet{i}.xml';root=ET.fromstring(z.read(n));cells={c.get('r'):c for c in root.iter(tag('c'))}
                for addr,val in zip(['B4','D4','F4'],choices[s['name']]):
                    c=cells[addr]
                    for el in list(c):c.remove(el)
                    c.set('t','inlineStr');is_=ET.SubElement(c,tag('is'));ET.SubElement(is_,tag('t')).text=val
                for c in cells.values():
                    if c.find(tag('f')) is not None:
                        v=c.find(tag('v'))
                        if v is not None:c.remove(v)
                repl[n]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        rewrite(p,repl);rec,_=lo(p,PRIVATE/f'filtro_{trial}','xlsx');w=openpyxl.load_workbook(rec,data_only=True)
        for name,values in choices.items():
            key='antiguas' if name.startswith('6_') else 'no_abiertas' if name.startswith('7_') else 'tendencia'
            dims=[3,4,5] if key=='no_abiertas' else [2,3,4]
            selected=[r for r in data[key] if all(v=='Todas' or r[j]==v for j,v in zip(dims,values))]
            s=w[name];checks={}
            if key=='antiguas':checks={'B8':len(selected),'B9':sum(r[10] for r in selected),'B10':sum(r[18]=='Sí' for r in selected),'B11':sum(r[15] for r in selected),'B12':sum(r[11] for r in selected),'B13':sum(r[17] for r in selected)}
            elif key=='no_abiertas':
                checks={'B8':len(selected),'B9':sum(r[9] for r in selected),'B10':sum(not r[9] for r in selected)}
                checks.update({f'P{8+i}':sum(r[0]==str(y) for r in selected) for i,y in enumerate(range(2010,2027))})
            else:
                checks={f'B{8+i}':sum(r[28]==cl for r in selected) for i,cl in enumerate(['ALZA','BAJA','ESTABLE','NUEVA','EN_EXTINCION','CERRADA'])};checks['B14']=len(selected)
                for start,rows in [('AH',sorted([r for r in selected if r[24]>0],key=lambda r:(-r[24],r[0]))),('AK',sorted([r for r in selected if r[24]<0],key=lambda r:(r[24],r[0])))]:
                    for j in range(10):checks[f'{start}{8+j}']=rows[j][0] if j<len(rows) else ''
            for cell,expected in checks.items():assert same(s[cell].value,expected),(trial,name,cell,s[cell].value,expected)
            results.append(dict(prueba=trial,hoja=name,seleccion=values,filas_seleccionadas=len(selected),controles=len(checks),estado='PASS'))
    return results

def main():
    fuentes();model=json.loads((OUT/'MODELO_AGREGADO.json').read_text());book=OUT/BOOK
    # No tablas solapadas y tres selectores por análisis.
    for s in model['sheets']:
        bounds=[]
        for b in s['blocks']:
            r=(b['row'],b['row']+len(b['rows']),b['column'],b['column']+len(b['headers'])-1)
            for o,id in bounds:assert r[1]<o[0] or r[0]>o[1] or r[3]<o[2] or r[2]>o[3],(s['name'],b['id'],id)
            bounds.append((r,b['id']))
    controls=json.loads((OUT/'CONTROL_CIFRAS.json').read_text());prior=json.loads((PRIVATE/OUT.name/'CONTROL_CIFRAS.json').read_text())
    index={(c['hoja'],c['celda']):c for c in controls}
    assert all(c==index[c['hoja'],c['celda']] for c in prior), 'Control previo modificado'
    q,fw=workbook_checks(book,PRIVATE/OUT.name/BOOK,controls,PRIVATE/'qa_resumen')
    assert len(fw.sheetnames)==9 and q['graficos']==4
    for name in ['6_Carreras_Antiguas','7_Ofertadas_No_Abiertas','8_Tendencia_Matricula']:assert len(fw[name].data_validations.dataValidation)==3
    q['controles_previos']=len(prior);q['filtros']=filters(book,model,json.loads((OUT/'DATOS_H006.json').read_text()))
    # Reproducción de modelos, no de vínculos ni análisis cerrados.
    files=['MODELO_AGREGADO.json','CONTROL_CIFRAS.json','RESULTADOS_H006.json','DATOS_H006.json'];hashes={n:sha(OUT/n) for n in files}
    subprocess.run([str(RUNTIME/'python/bin/python3'),str(OUT/'construir.py'),'--solo-modelo'],stdout=subprocess.DEVNULL,check=True);assert hashes=={n:sha(OUT/n) for n in files}
    q['reproduccion']=True
    setup_panels(book,model);pdf,_=lo(book,PRIVATE/'pdf_resumen','pdf');shutil.copy2(pdf,OUT/'VISTA_IMPRESION_H006.pdf');doc=pypdfium2.PdfDocument(pdf);rend=OUT/'renders_h006';rend.mkdir(exist_ok=True)
    q['pngs']=[]
    for i,p in enumerate(doc):
        dest=rend/f'panel_{i+1:02}.png';p.render(scale=1.5).to_pil().save(dest);q['pngs'].append(str(dest))
    ret=ROOT/'outputs/entrega_retencion_20261029';retbook=ret/'ENTREGA_RETENCION_IPSS_20261029.xlsx';rcontrols=json.loads((ret/'CONTROL_CELDAS.json').read_text());qr,rf=workbook_checks(retbook,PRIVATE/ret.name/retbook.name,rcontrols,PRIVATE/'qa_retencion')
    assert qr['controles']==1956 and qr['formulas']==41
    # Auditoría de privacidad de salidas agregadas, sin escanear fuentes individuales.
    pat=[r'\b\d{1,2}\.\d{3}\.\d{3}-[0-9kK]\b',r'(?<![A-Za-z0-9])\d{7,8}-[0-9kK](?![A-Za-z0-9])',r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}']
    for w in [fw,rf]:
        for s in w:
            for row in s:
                for c in row:
                    if c.data_type!='f' and isinstance(c.value,str):assert not any(re.search(p,c.value) for p in pat),(s.title,c.coordinate)
    initial=json.loads((PRIVATE/'ESTADO_INICIAL.json').read_text())
    allowed=['outputs/excel_resumen_cruce_20261030/','outputs/entrega_retencion_20261029/','docs/MAPA_RESULTADOS_GOBERNADOS.']
    for p,h in initial['hashes'].items():
        if not any(p.startswith(a) for a in allowed):assert sha(ROOT/p)==h,p
    assert subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT).decode().strip()==initial['head']
    current=subprocess.check_output(['git','status','--porcelain'],cwd=ROOT).decode().splitlines()
    clean=lambda ls:[s for s in ls if not any(a in s for a in allowed)]
    assert clean(current)==clean(initial['status'].splitlines())
    desktop=Path.home()/'Desktop/ENTREGA_RETENCION_20261031';desktop.mkdir(exist_ok=True)
    for b in [book,retbook]:shutil.copy2(b,desktop/b.name);assert sha(b)==sha(desktop/b.name)
    q.update(estado='PASS_TECNICO_PENDIENTE_APERTURA_VISUAL',privacidad='PASS',cambios_ajenos=24,sha256=sha(book),sha256_escritorio=sha(desktop/book.name),sin_commit_push=True)
    qr.update(sha256=sha(retbook),sha256_escritorio=sha(desktop/retbook.name));q['retencion']=qr
    dump(OUT/'QA_H006.json',q);print(json.dumps({k:v for k,v in q.items() if k not in ['filtros','pngs','retencion']},ensure_ascii=False))
if __name__=='__main__':main()
