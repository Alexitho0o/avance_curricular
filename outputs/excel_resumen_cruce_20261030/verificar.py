"""QA de presentación H005. openpyxl se usa solo para lectura y cotejo."""
from pathlib import Path
import json, re, zipfile, shutil, subprocess, math
import xml.etree.ElementTree as ET
import openpyxl, pypdfium2
from pypdf import PdfReader
from construir import OUT,ROOT,RUNTIME,BOOK,sha,dump,fuentes

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
ET.register_namespace('',NS)
def tag(s):return '{'+NS+'}'+s
def rewrite(p,repl):
    with zipfile.ZipFile(p) as z:content={n:z.read(n) for n in z.namelist()}
    content.update(repl);q=p.with_suffix('.tmp')
    with zipfile.ZipFile(q,'w',zipfile.ZIP_DEFLATED) as z:
        for n,v in content.items():z.writestr(n,v)
    q.replace(p)
def lo(p,work,fmt):
    profile=work/'perfil';(profile/'user').mkdir(parents=True,exist_ok=True)
    (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale" oor:op="fuse"><value>es-CL</value></prop><prop oor:name="ooLocale" oor:op="fuse"><value>es-CL</value></prop></item></oor:items>''')
    dest=work/fmt;dest.mkdir(exist_ok=True)
    r=subprocess.run([str(RUNTIME/'bin/override/soffice'),'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to',fmt,'--outdir',str(dest),str(p)],capture_output=True,text=True,timeout=120)
    target=dest/(p.stem+'.'+fmt)
    assert r.returncode==0 and target.exists(),r.stdout+r.stderr
    return target,r.stdout+r.stderr
def caches(p,recalc):
    replacements={}
    with zipfile.ZipFile(p) as a,zipfile.ZipFile(recalc) as b:
        for n in a.namelist():
            if not re.fullmatch(r'xl/worksheets/sheet\d+.xml',n):continue
            root=ET.fromstring(a.read(n));other=ET.fromstring(b.read(n));cells={c.attrib['r']:c for c in other.iter(tag('c'))}
            for c in root.iter(tag('c')):
                if c.find(tag('f')) is None:continue
                v=cells[c.attrib['r']].find(tag('v'));assert v is not None
                old=c.find(tag('v'))
                if old is not None:c.remove(old)
                c.append(ET.fromstring(ET.tostring(v)))
            replacements[n]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    rewrite(p,replacements)
def setup_panels(p,model):
    replacements={}
    with zipfile.ZipFile(p) as z:
        root=ET.fromstring(z.read('xl/workbook.xml'));dn=root.find(tag('definedNames'))
        if dn is None:
            dn=ET.Element(tag('definedNames'));calc=root.find(tag('calcPr'));root.insert(list(root).index(calc) if calc is not None else len(root),dn)
        for n in list(dn):
            if n.get('name')=='_xlnm.Print_Area':dn.remove(n)
        for i,s in enumerate(model['sheets']):
            a=ET.SubElement(dn,tag('definedName'),{'name':'_xlnm.Print_Area','localSheetId':str(i)})
            area=re.sub(r'([A-Z]+)([0-9]+)',r'$\1$\2',s['renders'][0])
            a.text=f"'{s['name']}'!{area}"
            n=f'xl/worksheets/sheet{i+1}.xml';sh=ET.fromstring(z.read(n));pr=sh.find(tag('sheetPr'))
            if pr is None:pr=ET.Element(tag('sheetPr'));sh.insert(0,pr)
            ET.SubElement(pr,tag('pageSetUpPr'),{'fitToPage':'1'})
            for old in ['pageMargins','pageSetup','printOptions']:
                el=sh.find(tag(old))
                if el is not None:sh.remove(el)
            idx=next((j for j,e in enumerate(sh) if e.tag in [tag('drawing'),tag('legacyDrawing'),tag('tableParts'),tag('extLst')]),len(sh))
            for el in [ET.Element(tag('printOptions'),{'headings':'0','gridLines':'0'}),ET.Element(tag('pageMargins'),{'left':'0.25','right':'0.25','top':'0.25','bottom':'0.25','header':'0','footer':'0'}),ET.Element(tag('pageSetup'),{'paperSize':'8','orientation':'landscape','fitToWidth':'1','fitToHeight':'1'})]:sh.insert(idx,el);idx+=1
            replacements[n]=ET.tostring(sh,encoding='utf-8',xml_declaration=True)
        replacements['xl/workbook.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    rewrite(p,replacements)
def main():
    sources=fuentes();book=OUT/BOOK;work=Path.home()/'.local/share/ipss/h005_20261030/qa';work.mkdir(parents=True,exist_ok=True)
    model=json.loads((OUT/'MODELO_AGREGADO.json').read_text());controls=json.loads((OUT/'CONTROL_CIFRAS.json').read_text())
    recalc,log=lo(book,work,'xlsx');fw=openpyxl.load_workbook(book,data_only=False);vw=openpyxl.load_workbook(recalc,data_only=True)
    differences=[];checked=[]
    for c in controls:
        actual=vw[c['hoja']][c['celda']].value;expected=c['esperado']
        same=math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12) if isinstance(actual,(int,float)) and isinstance(expected,(int,float)) else actual==expected
        if expected=='':same=actual in ['',None]
        checked.append(dict(c,obtenido=actual,coincide=same))
        if not same:differences.append({'hoja':c['hoja'],'celda':c['celda'],'esperado':expected,'obtenido':actual})
    dump(OUT/'CONTROL_RESULTADOS.json',checked)
    if differences:dump(OUT/'DIFERENCIAS_QA.json',differences);raise ValueError(str(differences[:10]))
    errors=[];formulas=0;hidden=[]
    for s in fw:
        if s.sheet_state!='visible':hidden.append(s.title)
        for row in s:
            for c in row:
                if c.data_type=='f':formulas+=1
                v=vw[s.title][c.coordinate]
                if v.data_type=='e':errors.append([s.title,c.coordinate,v.value])
    assert not errors and not hidden
    assert len(fw.sheetnames)==5 and sum(len(s._charts) for s in fw)==1
    assert len(fw['2_Programas_SIES']._charts[0].series)==2
    assert fw['2_Programas_SIES']._charts[0].grouping=='clustered'
    assert all(s.freeze_panes for s in fw)
    tables={s.title:len(s.tables) for s in fw};assert all(tables.values())
    for s in fw:
        for t in s.tables.values():assert t.autoFilter is not None
    caches(book,recalc)
    # Reproducibilidad de los datos de presentación, no de los productores.
    targets=['MODELO_AGREGADO.json','CIFRAS_CLAVE.json','CONTROL_CIFRAS.json','QA_FUENTES.json','CLASIFICACION_253_H005B.json','ATRIBUTOS_VERSIONES_H005B.json','COHERENCIA_TIPO_H005C.json','CLASIFICACION_253_H005C.json'];before={n:sha(OUT/n) for n in targets}
    subprocess.run([str(RUNTIME/'python/bin/python3'),str(OUT/'construir.py'),'--solo-modelo'],stdout=subprocess.DEVNULL,check=True)
    assert before=={n:sha(OUT/n) for n in targets}
    stats=json.loads((OUT/'CIFRAS_CLAVE.json').read_text())
    assert sum(stats['motivos'].values())==stats['universo_union']
    assert sum(stats['tipos_maestro'].values())==stats['planes']
    assert sum(stats['estados_2027'].values())==161
    assert sum(stats['historicos_por_nivel'].values())==stats['codigos_historicos']
    assert stats['enlazadas']==sum(r['esperado'] for r in controls if r['hoja']=='1_Base_Completa' and r['campo'].startswith('Enlazadas '))
    # Las 25.318 comprobaciones originales siguen vigentes aunque se haya
    # desplazado una tabla o renumerado una hoja. No se compara por posición.
    aliases={'3_Codigos_UMAS':'4_Codigos_UMAS','4_Cruce_SIES_UMAS':'5_Cruce_SIES_UMAS'}
    legacy=json.loads((OUT/'CONTROL_H005_ORIGINAL.json').read_text());assert len(legacy)==25318
    index={(c['hoja'],c['bloque'],c['registro'],c['campo']):c for c in controls}
    legacy_results=[]
    for old in legacy:
        key=(aliases.get(old['hoja'],old['hoja']),old['bloque'],old['registro'],old['campo'])
        new=index[key];actual=vw[new['hoja']][new['celda']].value;expected=old['esperado']
        same=math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-12) if isinstance(actual,(int,float)) and isinstance(expected,(int,float)) else actual==expected
        if expected=='':same=actual in ['',None]
        assert same,(key,expected,actual)
        legacy_results.append(dict(hoja=new['hoja'],celda=new['celda'],bloque=old['bloque'],registro=old['registro'],campo=old['campo'],coincide=True))
    dump(OUT/'PRESERVACION_H005.json',dict(controles=len(legacy),diferencias=0,sha256_base=sha(OUT/'CONTROL_H005_ORIGINAL.json'),controles_reubicados=legacy_results))
    newstats=stats['h005b']
    assert sum(newstats['oferta_por_nivel'].values())==newstats['en_oferta']
    assert sum(newstats['clasificacion_253'].values())==253
    assert sum(r['matriculas'] for r in newstats['clases_por_familia_anio'])==253
    # H005c: conservar los 34.510 controles anteriores salvo clasificación.
    baseline=Path.home()/'.local/share/ipss/h005c_20261030/respaldo_H005b'
    prev=json.loads((baseline/'CONTROL_CIFRAS.json').read_text());assert len(prev)==34510
    approved={'Resumen253','Clasificacion253','Evidencia253'};preserved=0;exceptions=[]
    for old in prev:
        key=(old['hoja'],old['bloque'],old['registro'],old['campo'])
        new=index.get(key)
        if old['bloque'] not in approved:
            assert new is not None and old['esperado']==new['esperado'] and old['formula']==new['formula'],key
            preserved+=1
        elif old['bloque']=='Evidencia253' and old['campo'] not in ['Clasificación','Clase','Fundamento / límite']:
            assert new is not None and old['esperado']==new['esperado'],key
            preserved+=1
        else:exceptions.append(key)
    co=json.loads((OUT/'COHERENCIA_TIPO_H005C.json').read_text())
    for r in co['anual']:
        assert r['correcto']+r['error']+r['no_evaluable']==r['enlazadas']
        assert r['enlazadas']==next(x[4] for x in stats['anual'] if x[0]==r['anio'])
        assert r['enlazadas']==sum(x['matriculas'] for x in co['cruce_tipos'] if x['anio']==r['anio'])
        assert r['error']==sum(x['matriculas'] for x in co['errores'] if x['anio']==r['anio'])
    # Diferencia exhaustiva de celdas contra el libro anterior (valores y fórmulas).
    bw=openpyxl.load_workbook(baseline/BOOK,data_only=False)
    bw_values=openpyxl.load_workbook(baseline/BOOK,data_only=True)
    cell_changes=[];added=[]
    for sw in fw:
        old=bw[sw.title]
        for row in sw:
            for cell in row:
                previous=old[cell.coordinate].value;current=cell.value
                old_value=bw_values[sw.title][cell.coordinate].value;new_value=vw[sw.title][cell.coordinate].value
                if previous==current and old_value==new_value:continue
                rec=dict(hoja=sw.title,celda=cell.coordinate,anterior=old_value,nuevo=new_value)
                if old[cell.coordinate].data_type=='f' or cell.data_type=='f':rec.update(formula_anterior=previous,formula_nueva=current)
                if previous is None:added.append(rec)
                else:
                    assert sw.title=='5_Cruce_SIES_UMAS' and 41<=cell.column<=60 and cell.row<=46,rec
                    cell_changes.append(rec)
        for row in old:
            for cell in row:
                if cell.value is not None and sw[cell.coordinate].value is None:
                    assert sw.title=='5_Cruce_SIES_UMAS' and 41<=cell.column<=60 and cell.row<=46
                    rec=dict(hoja=sw.title,celda=cell.coordinate,anterior=cell.value,nuevo=None)
                    if rec not in cell_changes:cell_changes.append(rec)
    dump(OUT/'CAMBIOS_CELDAS_H005C.json',dict(controles_previos=len(prev),conservados_sin_excepcion=preserved,controles_clasificacion_revisados=len(exceptions),celdas_modificadas=cell_changes,celdas_agregadas=added))
    lines=['# H005c · Cambios de celdas respecto de H005b','','Sólo cambia la clasificación; el universo y los campos originales de los 34 grupos se conservan. Las celdas añadidas están en el JSON.','','| Hoja | Celda | Valor anterior | Valor nuevo |','| --- | --- | --- | --- |']
    for x in cell_changes:
        def clean(v):return str(v if v is not None else 'vacío').replace('|','/').replace('\n',' / ')
        lines.append('| '+ ' | '.join(clean(x[k]) for k in ['hoja','celda','anterior','nuevo'])+' |')
    (OUT/'CAMBIOS_CELDAS_H005C.md').write_text('\n'.join(lines)+'\n')
    # Privacidad: esquema agregado; no columnas de identificación individual.
    bad=[];patterns=[r'\b\d{1,2}\.\d{3}\.\d{3}-[0-9kK]\b',r'(?<![A-Za-z0-9])\d{7,8}-[0-9kK](?![A-Za-z0-9])',r'[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}']
    for p in OUT.rglob('*'):
        if not p.is_file() or p.suffix not in ['.json','.md','.csv','.ndjson','.txt']:continue
        text=p.read_text(errors='replace')
        if any(re.search(x,text) for x in patterns):bad.append(p.name)
    for s in vw:
        for row in s:
            for c in row:
                if isinstance(c.value,str) and any(re.search(x,c.value) for x in patterns):bad.append(s.title+'!'+c.coordinate)
    assert not bad,bad
    init=json.loads((work.parent/'ESTADO_INICIAL.json').read_text());changed=[]
    for f,h in init['hashes'].items():
        if f in ['docs/MAPA_RESULTADOS_GOBERNADOS.md','docs/MAPA_RESULTADOS_GOBERNADOS.json']:continue
        if not (ROOT/f).is_file() or sha(ROOT/f)!=h:changed.append(f)
    assert not changed,changed
    assert subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT).decode().strip()==init['head']
    # PDF exclusivamente de paneles QA, conservando intacto el libro de entrega.
    panel=work/'PANELES_H005.xlsx';shutil.copy2(book,panel);setup_panels(panel,model);pdf,plog=lo(panel,work,'pdf')
    shutil.copy2(pdf,OUT/'PANELES_QA.pdf');doc=pypdfium2.PdfDocument(pdf);reader=PdfReader(pdf)
    rend=OUT/'renders_finales';rend.mkdir(exist_ok=True);pngs=[]
    assert len(doc)==5, 'Se requiere un panel final por hoja'
    for i,p in enumerate(doc):
        dest=rend/f'panel_{i+1:02}.png';p.render(scale=1.5).to_pil().save(dest);pngs.append(str(dest))
    text='\n'.join(p.extract_text() for p in reader.pages)
    assert '96,4' in text,'Porcentaje no visible en formato chileno'
    target=Path.home()/'Desktop/ENTREGA_RETENCION_20261031';target.mkdir(exist_ok=True);shutil.copy2(book,target/BOOK)
    assert sha(book)==sha(target/BOOK)
    qa=dict(estado='PASS_TECNICO; falta apertura y revisión visual final',fuentes=len(sources),controles=len(controls),controles_H005=25318,diferencias_H005=0,controles_nuevos=len(controls)-25318,diferencias=0,formulas=formulas,errores_formula=errors,
       tablas=tables,graficos=1,hojas_ocultas=0,privacidad=bad,esquema='Catálogo y conteos, sin lectura de microdatos ni enlaces nuevos',reproduccion_datos=True,
       archivos_previos_preservados=len(init['hashes']),sha256=sha(book),copia_escritorio=str(target/BOOK),mismo_hash=True,formato_porcentajes_es_CL=True,
       pngs=pngs,tipografia='Arial disponible; sustitución de Montserrat declarada',sin_commit_push=True)
    qa['H005c']=dict(controles_previos=34510,controles_no_reclasificados_conservados=preserved,controles_clasificacion_revisados=len(exceptions),celdas_modificadas=len(cell_changes),celdas_agregadas=len(added),coherencia_anual=co['anual'])
    qa['pngs_complementarios']=[str(p) for p in sorted((OUT/'renders').glob('*.png'))]
    dump(OUT/'QA.json',qa);(OUT/'RECALCULO.txt').write_text(log+'\n'+plog)
    print(json.dumps({k:v for k,v in qa.items() if k!='pngs'},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
