"""Lee respaldos de titulación; conserva un resultado no identificatorio en privado."""
import csv
import hashlib
import json
from pathlib import Path
import openpyxl

PRIVATE = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
OUT = Path('/Users/alexi/.local/share/ipss/busqueda_desempate_titulacion_20261026')
ROOT = Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Datos DAI (RAW) - 1.1 Datos crudos/1.2.1 Datos Institucionales/1.2.1.1 SIES/Enviados')
FOLDER = next(ROOT.glob('Titulaci*/2026'))
CAPTURE = next(FOLDER.glob('16683*.xlsx'))
LOAD = next(p for p in FOLDER.glob('*.xlsx') if p != CAPTURE)
CSV = next(p for p in FOLDER.glob('*.csv') if 'unificada' in p.name.lower())
TEAMS = next(Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Archivos de chat de Microsoft Teams').glob('1668382*.xlsx'))

def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for part in iter(lambda: stream.read(1048576), b''):
            h.update(part)
    return h.hexdigest()

def norm(v):
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v or '').strip().upper()

def body(v):
    return norm(v).replace('.', '').split('-')[0]

def records(sheet, max_rows=None):
    rows = sheet.iter_rows(max_row=max_rows, values_only=True)
    names = [str(x or '') for x in next(rows)]
    return names, [dict(zip(names, r)) for r in rows]

cases = json.loads(PRIVATE.read_text())
case_candidates = [c for c in cases if any(x['carrera'] == 'TCRE' for x in c['contextos_v11'])]
assert len(case_candidates) == 1
case = case_candidates[0]
ident, dv = case['RUT'].split('-')
initial = {str(p): digest(p) for p in [CAPTURE, LOAD, CSV, TEAMS]}
result = {'clasificacion': 'B fuentes, C contraste, E asignación', 'fuentes': [], 'observaciones': {}}
traces = []

def one_title(rows):
    return [r for r in rows if body(r.get('NUM_DOCUMENTO')) == ident and norm(r.get('DV')) == dv]

for book in [TEAMS, CAPTURE, LOAD]:
    w = openpyxl.load_workbook(book, read_only=True, data_only=False)
    summary = {'ruta': str(book), 'sha256': initial[str(book)], 'hojas': []}
    for sheet_name in ['Res_Tit_Unificada', 'Res_Tit_Carga']:
        if sheet_name not in w.sheetnames: continue
        s = w[sheet_name]
        # El rango límite evita el formato extendido de la copia original (1.048.576 filas).
        headers, rows = records(s, max_rows=500)
        matches = one_title(rows)
        item = {'hoja': sheet_name, 'filas_leidas_despues_encabezado': len(rows),
                'coincidencias_rut_dv': len(matches),
                'columnas_codcli_plan': [h for h in headers if 'CODCLI' in h.upper() or 'PLAN' in h.upper()],
                'columnas_puntaje_razones': [h for h in headers if 'SCORE' in h.upper() or 'REASON' in h.upper()]}
        if matches and 'TRACE_SCORE_CANDIDATO' in headers:
            r = matches[0]
            traces.append((norm(r['TRACE_CODCLI_DATOS_ALUMNOS']), norm(r['TRACE_SCORE_CANDIDATO']), norm(r['TRACE_REASONS_CANDIDATO'])))
            item['puntaje_guardado'] = r['TRACE_SCORE_CANDIDATO']
            item['codcli_traza_entre_expedientes'] = any(norm(r['TRACE_CODCLI_DATOS_ALUMNOS']) == norm(x['CODCLI']) for x in case['datosalumnos'])
            item['formula_puntaje'] = isinstance(r['TRACE_SCORE_CANDIDATO'], str) and r['TRACE_SCORE_CANDIDATO'].startswith('=')
            item['razones_guardadas'] = r['TRACE_REASONS_CANDIDATO']
        summary['hojas'].append(item)
    w.close()
    result['fuentes'].append(summary)

assert len(traces) == 2
result['observaciones']['copias_traza_mismo_candidato_puntaje_y_razones'] = traces[0] == traces[1]

raw = CSV.read_bytes()
for encoding in ['utf-8-sig', 'cp1252']:
    try: content = raw.decode(encoding); break
    except UnicodeDecodeError: pass
data = list(csv.reader(content.splitlines(), delimiter=';'))
matches = [r for r in data if len(r) >= 4 and body(r[2]) == ident and norm(r[3]) == dv]
result['fuentes'].append({'ruta': str(CSV), 'sha256': initial[str(CSV)], 'encoding_lectura': encoding,
                          'delimitador': ';', 'filas': len(data),
                          'distribucion_numero_columnas': {str(n): sum(len(r)==n for r in data) for n in sorted({len(r) for r in data})},
                          'coincidencias_campo_3_y_4_como_documento_dv': len(matches),
                          'advertencia': 'Posiciones contrastadas con el encabezado del XLSX; esto no demuestra recepción ni aceptación.'})
for p in [CAPTURE, LOAD, CSV, TEAMS]: assert digest(p) == initial[str(p)]
result['limitaciones'] = ['No se recupera el puntaje ni regla de desempate a partir de valores pegados.',
 'La coincidencia del documento no atribuye un plan o CODCLI a un título si no están declarados en el registro.',
 'No se deduce recepción oficial desde el nombre o la ubicación de un CSV.']
payload = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
for c in cases:
    assert c['RUT'].split('-')[0] not in payload
    for x in c['datosalumnos']: assert norm(x['CODCLI']) not in payload
OUT.mkdir(parents=True, mode=0o700, exist_ok=True)
target = OUT / 'RESULTADO_SIN_IDENTIFICADORES.json'
target.write_text(payload)
target.chmod(0o600)
print(json.dumps({'fuentes': len(result['fuentes']), 'hojas': [x['hojas'] for x in result['fuentes'] if 'hojas' in x],
                  'csv': result['fuentes'][-1], 'hashes_conservados': True}, ensure_ascii=False))
