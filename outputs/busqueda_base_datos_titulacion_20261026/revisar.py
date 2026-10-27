"""Revisa respaldos indicados por el usuario; información individual solo en memoria."""
import collections
import csv
import hashlib
import json
import subprocess
import unicodedata
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
PRIVATE = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
ROOT = next(p for p in Path('/Users/alexi/Library/CloudStorage').iterdir() if unicodedata.normalize('NFC', p.name) == 'OneDrive-InstitutoProfesionalSanSebastián') / '0 BASES DE DATOS'
DA = ROOT / 'base de datos 1900-2026.xlsx'
TIT = ROOT / 'TITULACION' / 'TITULADO_2025_web_27_05_2026_E.xlsx'
CSV = ROOT / 'TITULACION' / 'Titulados_IPSS_162_Consolidado_2007_2025.csv'

def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''): h.update(block)
    return h.hexdigest()

def txt(v):
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v or '').strip().upper()

def rut(v):
    return txt(v).replace('.', '').replace(' ', '')

cases = json.loads(PRIVATE.read_text())
matching = [c for c in cases if any(x['carrera'] == 'TCRE' for x in c['contextos_v11'])]
assert len(matching) == 1
case = matching[0]
expected_rut = rut(case['RUT'])
targets = {txt(x['CODCLI']) for x in case['datosalumnos'] if x['CODCLI']}
assert len(targets) == 3
hashes = {str(p): sha(p) for p in [DA, TIT, CSV]}
w = openpyxl.load_workbook(DA, read_only=True, data_only=True)
s = w['Hoja1']; rows = s.iter_rows(values_only=True); headers = [txt(x) for x in next(rows)]
indices = {k: headers.index(k) for k in ['CODCLI', 'RUT', 'CODCARPR', 'ESTADOACADEMICO', 'SITUACION',
                                          'ANOMATRICULA', 'PERIODOMATRICULA', 'ANOINGRESO', 'PERIODOINGRESO', 'JORNADA']}
reference = {txt(x['CODCLI']): x for x in case['datosalumnos']}
matches = []; n = 0
for row in rows:
    n += 1
    if txt(row[indices['CODCLI']]) in targets:
        ref = reference[txt(row[indices['CODCLI']])]
        pairs = [('CODCARPR', 'carrera'), ('ESTADOACADEMICO', 'estado'), ('SITUACION', 'situacion'),
                 ('ANOMATRICULA', 'ano_matricula'), ('PERIODOMATRICULA', 'periodo_matricula'),
                 ('ANOINGRESO', 'ano_ingreso'), ('PERIODOINGRESO', 'periodo_ingreso'), ('JORNADA', 'jornada')]
        matches.append({'rut_y_dv_iguales': rut(row[indices['RUT']]) == expected_rut,
                        'campos_revisados_coinciden': all(txt(row[indices[a]]) == txt(ref[b]) for a, b in pairs),
                        'carrera': txt(row[indices['CODCARPR']]),
                        'estado': txt(row[indices['ESTADOACADEMICO']]),
                        'situacion': txt(row[indices['SITUACION']])})
w.close()
da = {'ruta': str(DA), 'sha256': hashes[str(DA)], 'hoja': 'Hoja1', 'filas_datos': n,
      'columnas': len(headers), 'campos_plan': [x for x in headers if 'PLAN' in x],
      'campos_titulo_fecha': [x for x in headers if 'TITUL' in x],
      'codcli_objetivo': len(targets), 'filas_codcli_localizadas': len(matches),
      'filas_codcli_con_rut_dv': sum(x['rut_y_dv_iguales'] for x in matches),
      'filas_con_ocho_campos_iguales_al_respaldo_previo': sum(x['campos_revisados_coinciden'] for x in matches),
      'estados_observados': dict(collections.Counter(x['estado'] for x in matches)),
      'corte_extraccion': 'no acreditado'}
w = openpyxl.load_workbook(TIT, read_only=True, data_only=True)
t = w['in'];th = [txt(x) for x in next(t.iter_rows(values_only=True))]
tit = {'ruta': str(TIT), 'sha256': hashes[str(TIT)], 'hoja': 'in', 'columnas': len(th),
       'campos_codcli': [x for x in th if 'CODCLI' in x], 'campos_rut': [x for x in th if 'RUT' in x],
       'campos_total_titulaciones': [x for x in th if x == 'TOTAL TITULACIONES']}
w.close()
with CSV.open('rt', encoding='utf-8-sig', newline='') as f:
    ch = [txt(x) for x in next(csv.reader(f))]
csv_info = {'ruta': str(CSV), 'sha256': hashes[str(CSV)], 'columnas': len(ch),
            'contiene_mrun': 'MRUN' in ch, 'campos_codcli': [x for x in ch if 'CODCLI' in x],
            'campos_rut': [x for x in ch if x == 'RUT']}
for p in [DA, TIT, CSV]: assert sha(p) == hashes[str(p)]
cmd = ['rg', '--no-ignore', '--hidden', '-i', '-l']
for glob in ['*.py', '*.sql', '*.js', '*.mjs', '*.ipynb', '*.json', '*.md', '*.txt', '*.log']:
    cmd += ['-g', glob]
cmd += [r'TRACE_SCORE_CANDIDATO|TRACE_REASONS_CANDIDATO|Res_Tit_Unificada|\b2130\b', str(ROOT)]
proc = subprocess.run(cmd, capture_output=True, text=True, timeout=45)
assert proc.returncode in (0, 1) and not proc.stderr
result = {'clasificacion': 'B estructura y datos; C cruce; E vinculo y corte',
          'carpeta': str(ROOT), 'datos_alumnos': da, 'titulaciones_agregadas_2025': tit,
          'titulados_consolidado': csv_info,
          'busqueda_tecnica': {'extensiones': ['py', 'sql', 'js', 'mjs', 'ipynb', 'json', 'md', 'txt', 'log'],
                              'archivos_con_coincidencia_literal': len(proc.stdout.splitlines()),
                              'errores': bool(proc.stderr), 'codigo_salida': proc.returncode},
          'limitaciones': ['Sin fecha de extracción validada, no asignar estados a cortes históricos.',
                           'TIT y CSV no permiten atribuir título a CODCLI; no se realizó cruce RUT-MRUN.',
                           'Búsqueda técnica limitada a las extensiones listadas; no incluye comprimidos, Excel ni PDF.']}
payload = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
for c in cases:
    assert c['RUT'].split('-')[0] not in payload
    for d in c['datosalumnos']: assert txt(d['CODCLI']) not in payload
(OUT / 'RESULTADO_AGREGADO.json').write_text(payload)
print(json.dumps({'filas_datos_alumnos': n, 'columnas': len(headers),
                  'codcli_localizados': len(matches), 'codcli_con_rut_dv': da['filas_codcli_con_rut_dv'],
                  'estados': da['estados_observados'], 'plan_campos': da['campos_plan'],
                  'titulo_campos': da['campos_titulo_fecha'], 'mrun_en_consolidado': csv_info['contiene_mrun'],
                  'archivo_tecnico_coincidente': result['busqueda_tecnica']['archivos_con_coincidencia_literal'],
                  'hashes_conservados': True}, ensure_ascii=False))
