"""Evalúa cierre interno CO sin adjudicar códigos; conserva evidencia privada."""
import collections,datetime,hashlib,json,os,re
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).parent
PRIV=Path('/Users/alexi/.local/share/ipss/v11_fase2_20261026')
MU=ROOT/'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
OFERTA=Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/IPSS_CONCILIACION_UMAS_SIES_BASES/05_OFERTA_ACADEMICA_Y_DURACION/Oferta_Academica_2010_al_2026_SIES_05_06_2026_WEB_E.xlsx')
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
def text(v):
 if v is None:return ''
 if isinstance(v,float) and v.is_integer():return str(int(v))
 s=str(v).strip().upper();return '' if s in ('NULL','NONE','NAN') else s
def norm(n,d):
 n=text(n).replace('.','').replace(' ','');d=text(d)
 return str(int(n))+'-'+d if re.fullmatch(r'\d+',n) and re.fullmatch(r'[0-9K]',d) else None
def get(row,h,key):return row[h[key]] if key in h and h[key]<len(row) else None
assert sha(MU)=='19cf33ae4925a7c13ec4745394141a297f4f64be4052a240cb0d2c42ffeb53bd'
targets=json.loads((PRIV/'QA_RESPALDOS_HISTORICOS_DETALLE.json').read_text());by_person=collections.defaultdict(list)
for t in targets:by_person[t['persona']].append(t);t['mu_fuente_primaria']=[]
w=openpyxl.load_workbook(MU,read_only=True,data_only=True);s=w['BASE_RETENCION_MU'];it=s.iter_rows(values_only=True);h={text(v):i for i,v in enumerate(next(it))}
for i,row in enumerate(it,2):
 if text(get(row,h,'TIPO_DOC'))!='R':continue
 pe=norm(get(row,h,'N_DOC'),get(row,h,'DV'));tt=by_person.get(pe)
 if not tt:continue
 year=text(get(row,h,'ANIO_INFORMADO'));parts=[text(get(row,h,k)) for k in ['COD_SED','COD_CAR','JOR','VERSION']]
 if not all(p.isdigit() for p in parts):continue
 code='I162S%sC%sJ%sV%s'%tuple(parts)
 for t in tt:
  if year in {t['anio_T1'],'2026'}:t['mu_fuente_primaria'].append({'anio':year,'codigo':code,'fila':i,'vig':text(get(row,h,'VIG')),'ingreso_actual':text(get(row,h,'ANIO_ING_ACT')),'ingreso_origen':text(get(row,h,'ANIO_ING_ORI'))})
w.close()
families={'I162S3C114J2','I162S3C91J2','I162S3C115J2'};offer=collections.defaultdict(list)
w=openpyxl.load_workbook(OFERTA,read_only=True,data_only=True);s=w['in'];it=s.iter_rows(values_only=True);head=[str(x or '').strip() for x in next(it)];h={k:i for i,k in enumerate(head)}
for i,row in enumerate(it,2):
 if text(get(row,h,'Año')) not in {'2026','OFE_2026'}:continue
 code=text(get(row,h,'Código Único'));f=code.rsplit('V',1)[0]
 if f in families:offer[f].append({'codigo':code,'fila':i,'tipo':text(get(row,h,'Tipo Carrera')),'plan_especial':text(get(row,h,'Plan Especial')),'duracion':text(get(row,h,'Duración Total')),'vigencia':text(get(row,h,'Vigencia'))})
w.close()
co=[t for t in targets if t['co']];diff=[t for t in targets if t['diferencia_familia']];assert len(co)==73 and len(diff)==25
counts=collections.Counter(); later=collections.Counter()
for t in co:
 sameyear=[x for x in t['mu_fuente_primaria'] if x['anio']==t['anio_T1']];codes={x['codigo'] for x in sameyear};expected=t['codigo_T1']
 counts['MU_fuente_unica_igual_T1' if codes=={expected} else 'MU_fuente_ambigua_o_no_coincide']+=1
 counts['MU_fuente_vig1' if any(x['codigo']==expected and x['vig']=='1' for x in sameyear) else 'MU_fuente_sin_vig1']+=1
 j={'1':'D','2':'V','4':'O'}[t['jornada']]
 match=any((a:=k.split('|'))[1:3]==['2026','1'] and a[3]==t['plan'] and a[4]==t['carrera'] and a[5]==j for k in t['promedios'])
 counts['PROMEDIOS_T1_mismo_codcli_plan_carrera_jornada' if match else 'PROMEDIOS_T1_no_encontrado']+=1
 for label in ('DA_P2_2026','DA_P3_2026'):
  yes=any(x['sede']=='CO' and x['carrera']==t['carrera'] and x['jornada']==j for x in t['respaldos'][label]);counts[label+('_coincide' if yes else '_no_coincide')]+=1
 f=t['familias_destino'][0];codesoffer={x['codigo'] for x in offer[f]};counts['oferta_2026_solo_T1' if codesoffer=={expected} else 'oferta_2026_otras_o_sin_version']+=1
for t in diff:
 if t['anio']!='2025':continue
 m=[x for x in t['mu_fuente_primaria'] if x['anio']=='2026' and x['codigo'].rsplit('V',1)[0] in t['familias_destino']]
 if not m:continue
 later['contextos_mu2026_familia_destino']+=1
 for attr in ('ingreso_actual','ingreso_origen'):
  vals={x[attr] for x in m};numeric={int(x) for x in vals if x.isdigit()};
  if numeric and max(numeric)<=2025:later[attr+'_hasta_2025']+=1
  elif numeric:later[attr+'_posterior']+=1
  else:later[attr+'_sin_dato']+=1
 later['mu2026_codigo_unico' if len({x['codigo'] for x in m})==1 else 'mu2026_multiples_codigos']+=1
da={'clasificacion':'B fuentes; C validación; D autorización anterior; E propuesta pendiente','fecha_hora':datetime.datetime.now().astimezone().isoformat(),'co_73_controles':dict(counts),'oferta_2026':dict(offer),'mu_fuente_primaria_sha256':sha(MU),'oferta_2026_sha256':sha(OFERTA),'diferencias_2025_con_mu2026':dict(later),'alcance':'MU identifica persona, no CODCLI. Variante estricta T1 es control, no regla de adjudicación. No se modificaron bases ni denominadores.'}
file=PRIV/'EVALUACION_CIERRE_CO_V1_DETALLE_V2.json';fd=os.open(file,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
with os.fdopen(fd,'w') as f:json.dump(targets,f,ensure_ascii=False,default=str)
(OUT/'EVALUACION_CIERRE_CO_V1_AGREGADO_V2.json').write_text(json.dumps(da,ensure_ascii=False,indent=2));print(json.dumps(da,ensure_ascii=False))
