"""Inspecciona XLSX y listados de archivos comprimidos sin extraer datos personales."""
import collections
import hashlib
import json
import subprocess
import unicodedata
import zipfile
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
PRIVATE = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
BASE = next(p for p in Path('/Users/alexi/Library/CloudStorage').iterdir() if unicodedata.normalize('NFC', p.name) == 'OneDrive-InstitutoProfesionalSanSebastián') / '0 BASES DE DATOS'
def sha(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''): h.update(block)
    return h.hexdigest()
def txt(v):
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v or '').strip().upper()
def clean(v): return txt(v).replace('.', '').replace(' ', '')

cases = json.loads(PRIVATE.read_text())
case = next(c for c in cases if any(x['carrera'] == 'TCRE' for x in c['contextos_v11']))
targets = {txt(x['CODCLI']) for x in case['datosalumnos'] if x['CODCLI']}
assert len(targets) == 3
id_rut = clean(case['RUT'])
results = []; errors = []
files = sorted(BASE.rglob('*.xlsx'))
for path in files:
    try:
        digest = sha(path)
        content_hits = set()
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith(('xl/worksheets/', 'xl/sharedStrings', 'xl/workbook')) and n.endswith('.xml'):
                    with z.open(n) as stream:
                        overlap = b''
                        while True:
                            b = stream.read(1048576)
                            if not b: break
                            chunk = overlap + b
                            for word in [b'TRACE_SCORE_CANDIDATO', b'TRACE_REASONS_CANDIDATO', b'Res_Tit_Unificada']:
                                if word.lower() in chunk.lower(): content_hits.add(word.decode())
                            overlap = chunk[-100:]
        w = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for s in w:
            flags = []
            for row in s.iter_rows(min_row=1, max_row=min(12, s.max_row or 12), max_col=min(120,s.max_column or 120), values_only=True):
                values = {txt(v) for v in row if v is not None}
                if 'CODCLI' in values and ('PLAN_DE_ESTUDIO' in values or 'PLAN_ESTUDIOS' in values):
                    flags.append({'contiene_codigo_y_plan': True,
                                  'incluye_fecha_o_nombre_titulo': bool(values & {'FECHA_OBT_TIT_GRA','FECHA_TITULACION','TITULO','NOMBRE_TITULO','TITULO_OBTENIDO'})})
            sheets.append({'nombre': s.title, 'filas_declaradas': s.max_row, 'columnas': s.max_column,
                           'cabecera_codcli_y_plan_en_primeras_12_filas': flags})
        special = None
        if path.name == 'DATOSDEALUMNOS ORIGEN DE LOS DATOS2025.xlsx':
            s = w['BBDD']; it = s.iter_rows(values_only=True); header = [txt(x) for x in next(it)]
            idx = {k: header.index(k) for k in ['CODCLI', 'RUT', 'CODCARPR', 'ESTADOACADEMICO', 'SITUACION']}
            count = collections.Counter(); matched = set()
            for row in it:
                cod = txt(row[idx['CODCLI']])
                if cod not in targets: continue
                count['filas_codcli'] += 1; matched.add(cod)
                raw_rut = clean(row[idx['RUT']])
                if raw_rut == id_rut: count['rut_y_dv_iguales'] += 1
                elif raw_rut.split('-')[0] == id_rut.split('-')[0]: count['solo_cuerpo_rut_igual'] += 1
            special = {'hoja': 'BBDD', 'codcli_objetivo': len(targets),
                       'codcli_distintos_encontrados': len(matched), 'conteos': dict(count),
                       'corte_fuente': 'no acreditado por nombre'}
        w.close()
        assert sha(path) == digest
        results.append({'ruta': str(path), 'sha256': digest, 'hojas': sheets,
                        'etiquetas_traza_en_xml': sorted(content_hits), 'caso_matricula_2025': special})
    except Exception as exc:
        errors.append({'ruta': str(path), 'error': type(exc).__name__})

archives = []
for path in sorted([p for p in BASE.rglob('*') if p.is_file() and p.suffix.lower() in {'.zip','.rar'}]):
    try:
        digest = sha(path)
        cmd = subprocess.run(['/usr/bin/bsdtar','-tf',str(path)], capture_output=True, text=True, timeout=15)
        names = cmd.stdout.splitlines() if cmd.returncode == 0 else []
        extensions = collections.Counter(Path(n).suffix.lower() for n in names if n and not n.endswith('/'))
        candidates = sum(1 for n in names if any(s in n.casefold() for s in ['titulacion','titulación','res_tit_unificada','trace_score','codcli','migracion']))
        assert sha(path) == digest
        archives.append({'ruta': str(path), 'sha256': digest, 'listado_exitoso': cmd.returncode == 0,
                         'archivos': sum(extensions.values()), 'extensiones': dict(extensions),
                         'nombres_con_terminos_buscados': candidates})
    except Exception as exc:
        errors.append({'ruta': str(path), 'error': type(exc).__name__})
out = {'clasificacion': 'B estructura; C inspección; E atribución', 'raiz': str(BASE),
       'libros': results, 'comprimidos': archives, 'errores': errors,
       'limites': ['Búsqueda literal en XML, no prueba ausencia de código compilado o formulaciones distintas.',
                   'Se examinan solo primeras 12 filas por hoja para cabeceras; datos fuera de cabeceras pueden variar.',
                   'Archivos comprimidos examinados solo por listado y extensión; no se inspecciona contenido.',
                   'Las fechas de extracción y relación con título requieren documentación independiente.']}
payload = json.dumps(out, ensure_ascii=False, indent=2) + '\n'
for c in cases:
    assert c['RUT'].split('-')[0] not in payload
    for row in c['datosalumnos']: assert txt(row['CODCLI']) not in payload
(OUT / 'RESULTADO_AGREGADO.json').write_text(payload)
print(json.dumps({'libros': len(results), 'etiquetas_traza': sum(bool(x['etiquetas_traza_en_xml']) for x in results),
                  'hojas_con_codcli_plan': sum(bool(s['cabecera_codcli_y_plan_en_primeras_12_filas']) for x in results for s in x['hojas']),
                  'comprimidos': len(archives), 'listado_comprimidos_ok': sum(x['listado_exitoso'] for x in archives),
                  'errores': len(errors), 'resultado_caso': [x['caso_matricula_2025'] for x in results if x['caso_matricula_2025']]}, ensure_ascii=False))
