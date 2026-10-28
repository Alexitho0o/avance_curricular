"""Contraste agregado reproducible; no exporta identificadores personales."""
import csv
import hashlib
from collections import defaultdict, Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MU = ROOT / 'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
REF = ROOT / 'procesos/caracterizacion/gobernanza/sies/RETENCION/Retencion_ipss_2022_2024_GRANULAR.xlsx'
EXPECTED = ['19cf33ae4925a7c13ec4745394141a297f4f64be4052a240cb0d2c42ffeb53bd','e85d8c15c9fa2eccb25a376ea352f345ddcec91dbcf8176070b1ae3490c5d946']
for p, expected in zip((MU, REF), EXPECTED):
    with p.open('rb') as f:
        assert hashlib.file_digest(f, 'sha256').hexdigest() == expected, f'Hash no coincide: {p}'

def scalar(x):
    if x is None:
        return ''
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return str(x).strip().upper()

def key(row):
    return tuple(scalar(row.get(k)) for k in ('TIPO_DOC', 'N_DOC', 'DV'))

def code(row):
    parts = [scalar(row.get(k)) for k in ('COD_SED', 'COD_CAR', 'JOR', 'VERSION')]
    if not all(parts):
        return None
    return f'I162S{parts[0]}C{parts[1]}J{parts[2]}V{parts[3]}'

sheet = load_workbook(MU, read_only=True, data_only=True)['BASE_RETENCION_MU']
it = sheet.iter_rows(values_only=True)
header = next(it)
cohorts, observed = defaultdict(set), defaultdict(list)
missing_keys = missing_codes = 0
for values in it:
    row = dict(zip(header, values)); y = row.get('ANIO_INFORMADO')
    if y not in (2022, 2023, 2024, 2025):
        continue
    k = key(row)
    if not all(k):
        missing_keys += 1
        continue
    observed[(y, k)].append(row.get('VIG'))
    if y in (2022, 2023, 2024) and row.get('VIG') == 1 and row.get('ANIO_ING_ACT') == y:
        c = code(row)
        if c is None:
            missing_codes += 1
        else:
            cohorts[(y, c)].add(k)

sheet = load_workbook(REF, read_only=True, data_only=True)['Matriz Retencion Carrera']
it = sheet.iter_rows(values_only=True); header = next(it)
refs = {}
for values in it:
    row = dict(zip(header, values)); y = row['Ano de Cohorte']
    if y not in (2022, 2023, 2024):
        continue
    k = (y, scalar(row['Codigo Unico de Programa']))
    if k in refs:
        raise ValueError(f'Código de referencia duplicado: {k}')
    refs[k] = (row['Tamano de Cohorte'], row['Estudiantes Retenidos Institucional'])

result = []
non_vig_zero = set()
for y, c in sorted(set(cohorts) | set(refs)):
    people = cohorts.get((y, c), set())
    a = sum(any(v == 1 for v in observed.get((y + 1, k), ())) for k in people)
    b = sum((y + 1, k) in observed for k in people)
    for k in people:
        states = observed.get((y + 1, k), ())
        if states and 1 not in states and any(v != 0 for v in states):
            non_vig_zero.add((y, k))
    ref = refs.get((y, c))
    in_mu = (y, c) in cohorts
    if ref is None or not in_mu:
        status = 'sin referencia' if ref is None else 'sin cohorte MU'
    else:
        aa = (len(people), a) == ref
        bb = (len(people), b) == ref
        status = 'ambas' if aa and bb else 'solo A' if aa else 'solo B' if bb else 'ninguna'
    result.append([y, c, len(people) if in_mu else None, a if in_mu else None, b if in_mu else None,
                   b - a if in_mu else None, *(ref or (None, None)), status])

out = Path(__file__).parent / 'paridad_por_codigo.csv'
with out.open('w', encoding='utf-8', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['cohorte','codigo','cohorte_mu','retenidos_vig1','retenidos_cualquier_vig','solo_vig_no_1','cohorte_referencia','retenidos_referencia','resultado_paridad'])
    writer.writerows(result)
for y in (2022, 2023, 2024):
    r = [x for x in result if x[0] == y]
    print(y, dict(Counter(x[-1] for x in r)), 'solo_B_personas_suma_por_codigo', sum(x[5] for x in r if x[-1] == 'solo B'), 'solo_A_personas_suma_por_codigo', sum(x[5] for x in r if x[-1] == 'solo A'))
    print('  totales codigos mu', sum(x[2] or 0 for x in r), sum(x[3] or 0 for x in r), sum(x[4] or 0 for x in r), 'referencia',sum(x[6] or 0 for x in r),sum(x[7] or 0 for x in r))
print('llaves_incompletas',missing_keys,'codigos_incompletos_en_cohorte',missing_codes)
print('presentes_sin_vig1_y_con_estado_distinto_de_cero',len(non_vig_zero))
for y in (2022, 2023, 2024):
    by_person = Counter(k for (year, _), persons in cohorts.items() if year == y for k in persons)
    print(y, 'personas_con_mas_de_un_codigo',sum(n > 1 for n in by_person.values()))
