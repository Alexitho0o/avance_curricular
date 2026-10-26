"""Conserva XML de gráficos y hojas agregadas, reemplaza detalle personal por agregados."""
import zipfile,re,json,hashlib,xml.etree.ElementTree as E
from pathlib import Path
from xml.sax.saxutils import escape
import openpyxl
O=Path(__file__).parent;src=Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/retencion_puente_provisional_20261022_v6/RETENCION_IPSS_PROVISIONAL_PUENTE_20261022_V6.xlsx');dst=O/'RETENCION_IPSS_PROVISIONAL_PUENTE_20261022_V7.xlsx'
z=zipfile.ZipFile(src);new=zipfile.ZipFile(O/'suplemento.xlsx');payload=json.load(open(O/'payload.json'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda x:'{'+NS+'}'+x;E.register_namespace('',NS)
def strings(zz):
 if 'xl/sharedStrings.xml' not in zz.namelist():return []
 return [''.join(n.itertext()) for n in E.fromstring(zz.read('xl/sharedStrings.xml'))]
ss=strings(z);sn=strings(new)
# Añadir estilos nuevos sin alterar los originales.
base=z.read('xl/styles.xml').decode();b=E.fromstring(base);a=E.fromstring(new.read('xl/styles.xml'));offset={};num_map={}
for tag in ['numFmts','fonts','fills','borders','cellXfs']:
 old=b.find(q(tag));extra=a.find(q(tag));offset[tag]=len(old) if old is not None else 0
 if extra is None:continue
 if tag=='numFmts':
  nextid=max([163]+[int(n.get('numFmtId')) for n in old] if old is not None else [163])+1
  for n in extra:num_map[int(n.get('numFmtId'))]=nextid;n.set('numFmtId',str(nextid));nextid+=1
 if tag=='cellXfs':
  for n in extra:
   for attr,t in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:n.set(attr,str(int(n.get(attr,'0'))+offset[t]))
   n.set('numFmtId',str(num_map.get(int(n.get('numFmtId','0')),int(n.get('numFmtId','0')))));n.set('xfId','0')
 children=''.join(E.tostring(n,encoding='unicode').replace(' xmlns="'+NS+'"','') for n in extra)
 pat=rf'<{tag}\b[^>]*>.*?</{tag}>'
 found=re.search(pat,base,re.S)
 if found:
  chunk=found.group();chunk=re.sub(r'count="\d+"',f'count="{offset[tag]+len(extra)}"',chunk,count=1);chunk=chunk.replace(f'</{tag}>',children+f'</{tag}>');base=base[:found.start()]+chunk+base[found.end():]
 else:base=base.replace('<fonts',f'<{tag} count="{len(extra)}">{children}</{tag}><fonts',1)
wb=E.fromstring(z.read('xl/workbook.xml'));names=[s.get('name') for s in wb.find(q('sheets'))];excluded={p['replace_index']-1 for p in payload if 'replace_index' in p};affected={names[i] for i in excluded}
texts={name:z.read(f'xl/worksheets/sheet{i+1}.xml').decode() for i,name in enumerate(names)}
while True:
 before=set(affected)
 for name,txt in texts.items():
  fs=re.findall(r'<f\b[^>]*>(.*?)</f>',txt,re.S)
  if any(any(n in f for n in affected) for f in fs):affected.add(name)
 if before==affected:break
frozen=0
cellpat=r'<c\b[^>]*(?<!/)>.*?</c>'
def clean(txt,shared,freeze=False,addon=False):
 global frozen
 if addon:txt=txt.replace('<x:','<').replace('</x:','</').replace('xmlns:x=','xmlns=')
 def cell(m):
  global frozen
  s=m.group()
  if freeze and re.search(r'<f(?:\s|>)',s):
   assert '<v' in s,'Formula sin cache';s=re.sub(r'<f\b[^>]*>.*?</f>|<f\b[^>]*/>','',s,flags=re.S);frozen+=1
  if re.search(r'\bt="s"',s):
   idx=int(re.search(r'<v>(\d+)</v>',s).group(1));text=shared[idx]
   if 'fórmulas en vivo' in text:text='C · Foto de valores históricos preservados; detalle individual excluido. No recalcula con filtros en esta copia agregada.'
   s=re.sub(r'\bt="s"','t="inlineStr"',s);s=re.sub(r'<v>.*?</v>','<is><t xml:space="preserve">'+escape(text)+'</t></is>',s,flags=re.S)
  if addon:
   if re.search(r'\bs="\d+"',s):s=re.sub(r'\bs="(\d+)"',lambda m:f's="{int(m.group(1))+offset["cellXfs"]}"',s,count=1)
   else:s=s.replace('<c ','<c s="'+str(offset['cellXfs'])+'" ',1)
  return s
 txt=re.sub(cellpat,cell,txt,flags=re.S)
 if freeze:txt=re.sub(r'<dataValidations\b.*?</dataValidations>','',txt,flags=re.S)
 return txt
repl={};new_names=[]
for i,p in enumerate(payload,1):
 t=clean(new.read(f'xl/worksheets/sheet{i}.xml').decode(),sn,addon=True)
 if 'replace_index' in p:repl[f'xl/worksheets/sheet{p["replace_index"]}.xml']=t
 else:new_names.append(p['name']);repl[f'xl/worksheets/sheet{len(names)+len(new_names)}.xml']=t
# No datos individuales permanecen en sharedStrings; no se copian relaciones de hojas sustituidas.
skip={'xl/sharedStrings.xml','xl/calcChain.xml'}
for i in excluded:
 rel=f'xl/worksheets/_rels/sheet{i+1}.xml.rels'
 if rel in z.namelist():
  for n in E.fromstring(z.read(rel)):
   target=n.get('Target','')
   if 'table' in target:skip.add('xl/'+target.replace('../','').lstrip('/').removeprefix('xl/'))
  skip.add(rel)
work=z.read('xl/workbook.xml').decode();rels=z.read('xl/_rels/workbook.xml.rels').decode();ct=z.read('[Content_Types].xml').decode()
for j,name in enumerate(new_names,len(names)+1):
 work=work.replace('</sheets>',f'<sheet name="{escape(name)}" sheetId="{j}" r:id="rIdProvisional{j}"/></sheets>')
 rels=rels.replace('</Relationships>',f'<Relationship Id="rIdProvisional{j}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{j}.xml"/></Relationships>')
 ct=ct.replace('</Types>',f'<Override PartName="/xl/worksheets/sheet{j}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
rels=re.sub(r'<Relationship\b[^>]*(?:sharedStrings|calcChain)[^>]*/>','',rels)
for name in skip:ct=re.sub(r'<Override\b[^>]*PartName="/'+re.escape(name)+r'"[^>]*/>','',ct)
repl.update({'xl/styles.xml':base,'xl/workbook.xml':work,'xl/_rels/workbook.xml.rels':rels,'[Content_Types].xml':ct})
with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
 for name in z.namelist():
  if name in skip:continue
  data=z.read(name)
  if name in repl:data=repl.pop(name).encode()
  elif re.fullmatch(r'xl/worksheets/sheet\d+\.xml',name):
   i=int(re.search(r'sheet(\d+)',name).group(1))-1;data=clean(data.decode(),ss,freeze=names[i] in affected).encode()
  out.writestr(name,data)
 for name,data in repl.items():out.writestr(name,data.encode())
# Verificación independiente, valores históricos y gráficos.
oldw=openpyxl.load_workbook(src,read_only=True,data_only=True);outw=openpyxl.load_workbook(dst,read_only=True,data_only=True);wf=openpyxl.load_workbook(dst,read_only=True,data_only=False)
changes=[];errors=[]
for i,name in enumerate(names):
 if i in excluded:continue
 old=oldw[name];cur={c.coordinate:c.value for row in outw[name] for c in row if c.value is not None}
 for row in old:
  for c in row:
   if c.value is not None and isinstance(c.value,(int,float)) and cur.get(c.coordinate)!=c.value:changes.append([name,c.coordinate])
for s in outw:
 for row in s:
  for c in row:
   if c.data_type=='e':errors.append([s.title,c.coordinate,c.value])
zz=zipfile.ZipFile(dst);charts=[n for n in z.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
qa={'fuente':str(src),'sha256_fuente':hashlib.sha256(src.read_bytes()).hexdigest(),'salida':str(dst.resolve()),'sha256':hashlib.sha256(dst.read_bytes()).hexdigest(),'hojas':len(outw.sheetnames),'graficos':len(charts),'graficos_xml_identicos':all(z.read(n)==zz.read(n) for n in charts),'cambios_numericos_historicos':changes,'errores_celdas':errors,'formulas_dependientes_fijadas':frozen,'hojas_con_historicos_fijados_en_esta_ejecucion':sorted(affected) if frozen else [],'hojas_actualizadas':[names[i] for i in sorted(excluded)],'detalle_personal_excluido_en_v1':['9_Base_Analitica_Detalle','11_BASE_RETENCION_MU_RAW','20_Foto_BBDD','24_Foto_Corte_T3_2025','25_Foto_Corte_T3_2026'],'estado':'PENDIENTE_APROBACION_GOBERNANZA'}
assert not changes and not errors and qa['graficos_xml_identicos']
s=outw['29_Actualizacion_Provisional'];assert abs(s['G5'].value-24/63)<1e-12 and abs(s['G7'].value-8/63)<1e-12 and abs(s['G8'].value-19/63)<1e-12 and s['H8'].value is None and s['H5'].value==0
s=outw['31_Comparacion_Historica'];qa['desfases_previos_tasa_fuente_vs_cociente']=[]
for row in s.iter_rows(min_row=5,values_only=True):
 if row[0] is None:continue
 assert row[2:5]==row[5:8]
 if row[3] and isinstance(row[4],(int,float)) and abs(row[4]-row[2]/row[3])>1e-10:qa['desfases_previos_tasa_fuente_vs_cociente'].append({'grupo':row[0],'ano':row[1],'numerador':row[2],'denominador':row[3],'tasa_fuente':row[4],'tasa_cociente':row[2]/row[3]})
(O/'QA_ENTREGA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa,ensure_ascii=False))
