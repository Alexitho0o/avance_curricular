#!/usr/bin/env python3
"""H002f: catálogo anual agregado; originales y atribuciones personales intactos."""
import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import sys

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DATA = ROOT / 'data/paridad_codigo_sies/anual_20261029'
PRIVATE = Path.home() / '.local/share/ipss/h002f_20261029'
sys.path.insert(0, str(ROOT))
from common.analisis_desercion.h001 import normalizar as txt, codigo
from common.diccionario_sies_seguimiento import rut_normalizado


def sha(p):
    h = hashlib.sha256()
    with Path(p).open('rb') as f:
        for b in iter(lambda: f.read(1048576), b''):
            h.update(b)
    return h.hexdigest()


def js(x):
    return json.dumps(x, ensure_ascii=False, sort_keys=True, indent=2) + '\n'


def read(p):
    return json.loads(Path(p).read_text())


def ident(tipo, doc, dv):
    """Tipo explícito; para R reutiliza normalizador gobernado, sin fabricar DV."""
    t, d, v = txt(tipo).upper(), txt(doc), txt(dv).upper()
    if not t or not d or (t == 'R' and not v):
        return None
    if t == 'R':
        try:
            b, v = rut_normalizado(d + '-' + v).split('-')
            return (t, b, v)
        except ValueError:
            return None
    return (t, d, v)


def umas_ident(r):
    s = txt(r.get('RUT'))
    if '-' in s:
        try:
            b, d = rut_normalizado(s).split('-')
            if 'DIG' in r and txt(r['DIG']) and txt(r['DIG']).upper() != d:
                return None
            return ident('R', b, d)
        except ValueError:
            return None
    if 'DIG' in r:
        return ident('R', s, r['DIG'])
    return None


def number(v):
    s = txt(v)
    return int(float(s)) if s and re.fullmatch(r'\d+(?:\.0+)?', s) else None


def family(code):
    return re.sub(r'V\d+$', '', code)


def career(code):
    m = re.fullmatch(r'I162S\d+C(\d+)J\d+V\d+', code)
    return m[1] if m else code


def selected(year, vig):
    # Patrón observado B/C del puente; NO regla normativa transversal de VIG.
    return year == 2025 or txt(vig) == '1'


def historical_state(n, total):
    if total <= 0:
        return 'NO_CUADRA'
    if n == total:
        return 'CUADRA_TOTAL'
    return 'APROXIMA' if abs(n-total) / total <= .10 else 'NO_CUADRA'


def classify_group(codes, balances):
    if len(codes) > 1:
        return 'REPARTIDO'
    return 'CUADRA_UNICO' if codes and all(balances[c] for c in codes) else 'UNICO_SIN_CUADRATURA'


def choose(candidates):
    """No desempatar un expediente/plan por orden, mayoría o igualdad de totales."""
    cc = set(candidates)
    if len({x[0] for x in cc}) > 1:
        return None, 'VARIOS_EXPEDIENTES_COMPATIBLES'
    if len(cc) > 1:
        return None, 'VARIOS_PLANES_O_CONTEXTOS_ANUALES'
    return (next(iter(cc)), '') if cc else (None, 'SIN_CONTEXTO_COMPATIBLE')


def sheet_rows(path, sheet, fields):
    import openpyxl
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        s = w[sheet]; it = s.iter_rows(values_only=True)
        header = [txt(v) for v in next(it)]
        idx = {k: [i for i, v in enumerate(header) if v == k] for k in fields}
        for line, rr in enumerate(it, 2):
            d = {}
            for k, indices in idx.items():
                values = {txt(rr[i]) for i in indices if i < len(rr) and txt(rr[i])}
                d[k] = next(iter(values)) if len(values) == 1 else ''
            yield line, d
    finally:
        w.close()


def load_inputs():
    sources = read(OUT/'FUENTES.json')
    for e in sources:
        if sha(e['ruta']) != e['sha256']:
            raise ValueError('Fuente alterada: ' + e['id'])
    src = {e['id']: Path(e['ruta']) for e in sources}
    cat = read(src['data/catalogo_fuentes/CATALOGO_FUENTES.json'])['catalogo']
    mu_entry = next(e for e in cat if e['sha256'].startswith('19cf33ae'))
    mu_fields = ['ANIO_INFORMADO','TIPO_DOC','N_DOC','DV','COD_SED','COD_CAR','JOR','VERSION','VIG']
    mu = [(line, r) for line, r in sheet_rows(src[mu_entry['archivo']], 'BASE_RETENCION_MU', mu_fields)]
    # Proyección mínima: no se mantienen nombres, contactos, comentarios ni notas.
    da = []; coverage = []; skipped = []
    da_fields = ['CODCLI','CODCARPR','RUT','SEDE','JORNADA','ANOINGRESO','ANOMATRICULA']
    import openpyxl
    for name, p in sorted(src.items()):
        if not name.startswith('DATOSDEALUMNOS'):
            continue
        w = openpyxl.load_workbook(p, read_only=True, data_only=True)
        usable = []
        for sn in w.sheetnames:
            if sn != 'Hoja1':
                continue
            h = [txt(v) for v in next(w[sn].iter_rows(values_only=True))]
            if {'CODCLI','RUT','CODCARPR'}.issubset(h):
                usable.append(sn)
        w.close()
        if not usable:
            skipped.append({'fuente':name,'motivo':'Sin Hoja1 con encabezado CODCLI/RUT/CODCARPR; no inferir estructura de pivote ni identificador.'})
            continue
        for sn in usable:
            n = 0
            for line, r in sheet_rows(p, sn, da_fields):
                if r['CODCLI'] and r['CODCARPR']:
                    da.append((name, sn, line, r)); n += 1
            coverage.append({'fuente':name,'hoja':sn,'filas':n,'uso':'identidad, expediente y extremos de año; no vigencia histórica'})
    pp = []
    fields = ['CODCLI','RUT','DIG','CODCARR','PLAN_DE_ESTUDIO','JORNADA','ANO','PERIODO']
    for name, p in sorted(src.items()):
        if not name.startswith('PROMEDIOSDEALUMNOS'):
            continue
        n = 0
        for line, r in sheet_rows(p, 'Hoja1', fields):
            pp.append((name, line, r)); n += 1
        coverage.append({'fuente':name,'hoja':'Hoja1','filas':n,'uso':'actividad anual; unión de expedientes, no suma de asignaturas'})
        print('Leído', name, n, 'filas; salida solo agregada', flush=True)
    return sources, src, mu, da, pp, coverage, skipped


def build():
    sources, src, mu, da, pp, coverage, skipped = load_inputs()
    table = read(src['data/paridad_codigo_sies/20261029/TABLA_PARIDAD_2026_2027.json'])['filas']
    groups = read(src['data/versiones_por_cohorte/20261028/grupos_existentes.json'])
    master = read(src['data/versiones_historicas_sies_umas/fase_5_20261015/maestro.json'])
    profiles = read(src['outputs/paridad_codigo_sies_20261029/PERFIL_PLANES_H002C.json'])['perfiles']
    fam_groups = defaultdict(set); plan_sites = defaultdict(set)
    for g in groups:
        for f in g['familias']:
            fam_groups[f].add((g['codcarpr'],g['sede'],g['jornada']))
    for g in master['equivalencias_programa_2026']:
        fam_groups[family(g['codigo_unico'])].add((g['codcarpr'],g['sede_umas'],g['jornada_umas']))
    for p in profiles:
        plan_sites[(p['codcarpr'],p['CODPESTUD'])].add(p['sede_umas'])
    for p in master['planes_umas']:
        for ctx in p['contextos_catalogo']:
            plan_sites[(ctx['Codcarr'],p['CODPESTUD'])].add(ctx['Sede'])
    public = {}; public_lines = {}; nonstandard = []
    with src['Matricula_2007_2026_WEB_10_07_2026.csv'].open(encoding='latin1',newline='') as f:
        for line, r in enumerate(csv.DictReader(f,delimiter=';'), 2):
            if r['CÓDIGO DE INSTITUCIÓN'] != '162': continue
            y = int(r['AÑO'].removeprefix('MAT_')); c = r['CÓDIGO CARRERA']; k = (y,c)
            if k in public: raise ValueError('Código público anual repetido')
            public[k] = {'total':number(r['TOTAL MATRÍCULA']), 'primer_anio':number(r['TOTAL MATRÍCULA PRIMER AÑO']),
                         'nombre':r['NOMBRE CARRERA'], 'nivel':r['NIVEL GLOBAL']}
            public_lines[k] = line
            if not re.fullmatch(r'I162S\d+C\d+J\d+V\d+',c): nonstandard.append({'anio':y,'codigo':c})
    # Índice interno por identidad exacta; una foto actual no es una serie anual.
    da_index = defaultdict(set); owners = defaultdict(set); endpoints = defaultdict(set)
    da_sites = defaultdict(set); da_refs = defaultdict(list); first_umas = defaultdict(set)
    for name, sn, line, r in da:
        pe = umas_ident(r); cli = r['CODCLI']; car = r['CODCARPR']
        if not pe: continue
        owners[cli].add(pe); da_index[pe].add((cli,car))
        if r['SEDE']: da_sites[(pe,cli,car)].add(r['SEDE'])
        da_refs[(pe,cli,car)].append((name,sn,line))
        for field in ['ANOINGRESO','ANOMATRICULA']:
            y = number(r[field])
            if y and 1900 < y <= 2026:
                endpoints[(car,r['SEDE'],r['JORNADA'],y,field)].add(cli)
                first_umas[car].add(y)
    activity = defaultdict(set); act_refs = defaultdict(list); historical = defaultdict(set)
    activity_diagnostics = Counter(); periods = Counter()
    for name, line, r in pp:
        y = number(r['ANO']); pe = umas_ident(r); cli = r['CODCLI']; car = r['CODCARR']; plan = r['PLAN_DE_ESTUDIO']; j = r['JORNADA']
        if not pe or not cli or not y or not (1900 < y <= 2026):
            activity_diagnostics['sin_identidad_expediente_anio'] += 1; continue
        owners[cli].add(pe); periods[r['PERIODO']] += 1
        if car.upper() in {'','NULL','NAN','NONE'} or plan.upper() in {'','NULL','NAN','NONE'}:
            activity_diagnostics['sin_carrera_o_plan'] += 1; continue
        first_umas[car].add(y)
        sites = plan_sites[(car,plan)]
        # Única sede de catálogo y ninguna sede DA contradictoria para el expediente.
        observed = da_sites[(pe,cli,car)]
        if len(sites) != 1 or (observed and not observed <= sites):
            activity_diagnostics['sede_catalogo_ausente_multiple_o_conflicto_DA'] += 1; continue
        site = next(iter(sites)); ctx = (cli,car,plan,site,j)
        activity[(pe,y)].add(ctx)
        act_refs[(pe,y,ctx)].append((name,'Hoja1',line,r['PERIODO']))
        historical[(car,plan,site,j,y)].add(cli)
    # Totales directos desde MU, sin leer valores de H002 para contar o asignar.
    muc = Counter(); muc_all = Counter(); vigs = defaultdict(Counter)
    for _,r in mu:
        y = number(r['ANIO_INFORMADO']); c = codigo(r)
        if y not in range(2022,2027): continue
        muc_all[(y,c)] += 1; vigs[(y,c)][r['VIG']] += 1
        if selected(y,r['VIG']): muc[(y,c)] += 1
    balances = []; bindex = {}
    for y,c in sorted(set(muc_all) | {k for k in public if k[0]>=2022}):
        p = public.get((y,c)); total = p['total'] if p else 0
        fuera = bool(p and p['nivel'] != 'Pregrado' and not muc_all[(y,c)])
        delta = total - muc[(y,c)]
        entry = {'anio':y,'codigo_sies':c,'nivel_publicacion':p['nivel'] if p else 'NO_PUBLICADO',
                 'filas_mu':muc_all[(y,c)],'por_vig':dict(sorted(vigs[(y,c)].items())),
                 'mu_seleccionada':muc[(y,c)],'sies_total':total,'diferencia_sies_menos_mu':delta,
                 'regla_vig':'TODAS_OBSERVADO_PRECARGA_AC2026' if y==2025 else 'VIG_1_PATRON_PUENTE',
                 'estado':'FUERA_ALCANCE_POSTITULO' if fuera else ('CUADRA' if delta==0 else 'DIFERENCIA_E'),
                 'ancla_codigo':not fuera and delta==0 and total>0 and y!=2026,
                 'fila_publicacion':public_lines.get((y,c)), 'clase':'B/C' if delta==0 else 'B/C/E'}
        balances.append(entry); bindex[(y,c)] = entry
    # Verifica el puente existente independientemente; no ajusta ninguna cifra.
    bridge_ok = True
    with src['procesos/caracterizacion/gobernanza/sies/TABLA_PUENTE_CSV_BASE_RETENCION_MU.csv'].open(encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            k = (int(r['ANIO']),r['CODIGO_CARRERA_CSV'])
            bridge_ok &= (muc_all[k]==number(r['BASE_TOTAL_FILAS']) and vigs[k]['1']==number(r['BASE_VIG1'])
                          and public.get(k,{'total':0})['total']==number(r['TOTAL_MATRICULA_CSV']))
    if not bridge_ok: raise ValueError('Totales directos no reproducen el puente gobernado')
    counts = Counter(); details = []; selected_people = defaultdict(set); linked_people = defaultdict(set)
    for line,r in mu:
        y = number(r['ANIO_INFORMADO'])
        if y not in range(2022,2027): continue
        c = codigo(r); pe = ident(r['TIPO_DOC'],r['N_DOC'],r['DV']); include = selected(y,r['VIG'])
        allowed = fam_groups[family(c)]; possible = []
        if not pe: reason = 'IDENTIDAD_NO_COMPARABLE'
        elif not allowed: reason = 'FAMILIA_SIN_MAPEO'
        elif not da_index[pe]: reason = 'SIN_IDENTIDAD_EXACTA_DATOS_ALUMNOS'
        else:
            for ctx in activity[(pe,y)]:
                cli,car,plan,site,j = ctx
                if (car,site,j) in allowed and (cli,car) in da_index[pe] and owners[cli]=={pe}:
                    possible.append(ctx)
            reason = 'SIN_ACTIVIDAD_ANUAL_COMPATIBLE'
        chosen, why = choose(possible)
        if possible and not chosen: reason = why
        if chosen:
            cli,car,plan,site,j = chosen; reason = 'ENLAZADO_UNICO'
        else:
            cli='';car='SIN_ENLACE';plan='SIN_ENLACE';site='';j=''
        counts[(y,c,car,plan,site,j,reason,include)] += 1
        if include:
            selected_people[y].add(pe)
            if chosen: linked_people[y].add(pe)
        details.append({'fila_mu':line,'anio':y,'codigo_sies':c,'identidad':pe,'vig':r['VIG'],
                        'incluida':include,'motivo':reason,'seleccion':chosen,
                        'candidatos':sorted(possible),'referencias_DA':da_refs[(pe,cli,car)] if chosen else [],
                        'expedientes_DA_familia':sorted((cl,ca) for cl,ca in da_index[pe] if ca in {z[0] for z in allowed}),
                        'actividad_anual_persona':sorted(activity[(pe,y)]),
                        'referencias_PROMEDIOS':act_refs[(pe,y,chosen)] if chosen else []})
    PRIVATE.mkdir(parents=True,exist_ok=True); PRIVATE.chmod(0o700)
    private_path = PRIVATE/'ENLACES_MU_PLAN.json'; private_path.write_text(js(details)); private_path.chmod(0o600)
    # Clasificación sobre incluidos; excluidos VIG se conservan en columnas separadas.
    allkeys = {k[:7] for k in counts}; groups_codes = defaultdict(set)
    for (y,c,car,plan,site,j,reason,inc),n in counts.items():
        if inc and reason=='ENLAZADO_UNICO': groups_codes[(y,car,plan,site,j)].add(c)
    matrix=[]
    for y,c,car,plan,site,j,reason in sorted(allkeys):
        incl=counts[(y,c,car,plan,site,j,reason,True)]; excl=counts[(y,c,car,plan,site,j,reason,False)]
        codes=groups_codes[(y,car,plan,site,j)]; bal=bindex[(y,c)]
        state = 'SIN_ENLACE' if reason!='ENLAZADO_UNICO' else (
            classify_group(codes,{cc:bindex[(y,cc)]['estado']=='CUADRA' for cc in codes}) if incl else 'FUERA_SELECCION_VIG')
        matrix.append({'anio':y,'codigo_sies':c,'codcarpr':car,'plan_umas':plan,'sede_umas':site,'jornada_umas':j,
                       'matriculas_mu_incluidas':incl,'matriculas_mu_excluidas_vig':excl,'mu_codigo_total':bal['mu_seleccionada'],
                       'sies_codigo_total':bal['sies_total'],'sin_enlace_motivo':reason if reason!='ENLAZADO_UNICO' else '',
                       'estado':state,'ancla_codigo':bal['ancla_codigo'],'clase':'C' if state=='CUADRA_UNICO' else 'C/E',
                       'unidad':'matrículas MU','sede_respaldo':'Catálogo y DA concordantes; no certifica sede individual al corte' if site else '',
                       'vigencia_historica':'Actividad anual observada; no estado actual ni foto al 30-abr',
                       'fuente_conteo':'BASE_MU_142342 + DATOSALUMNOS + PROMEDIOS; FUENTES.json',
                       'control_codigo':bal['estado'],'codigo_es_destino_observado':True})
    # Etapa anterior a 2022: cada candidato se compara por separado, no se reparte SIES entre planes.
    hist=[]; first_sies=defaultdict(set)
    for (y,c),p in sorted(public.items()):
        if p['total'] and p['total']>0: first_sies[career(c)].add(y)
        if y>=2022 or not p['total']: continue
        cg=fam_groups[family(c)]; candidates=[]
        for (car,plan,site,j,ay),ids in sorted(historical.items()):
            if ay==y and (car,site,j) in cg: candidates.append((car,plan,site,j,len(ids)))
        if not candidates: candidates=[('SIN_ENLACE','SIN_ENLACE','','',0)]
        for car,plan,site,j,n in candidates:
            ep={f:len(endpoints[(car,site,j,y,f)]) for f in ['ANOINGRESO','ANOMATRICULA']}
            state=historical_state(n,p['total'])
            row={'anio':y,'codigo_sies':c,'codcarpr':car,'plan_umas':plan,'sede_umas':site,'jornada_umas':j,
                 'expedientes_actividad_anual':n,'sies_codigo_total':p['total'],'diferencia_interno_menos_sies':n-p['total'],
                 'diferencia_porcentaje':round(100*(n-p['total'])/p['total'],4),'estado':state,'clase':'C/E',
                 'unidad':'CODCLI distintos con actividad versus matrículas SIES (universos distintos)',
                 'candidatos_comparados_mismo_codigo':len(candidates),'conteos_DA_extremos_no_interpolar':ep,
                 'asignacion_version_autorizada':False,'ancla_codigo':False,'codigo_es_destino_observado':False,
                 'alcance':'Candidato de catálogo, sin vigencia temporal acreditada; igualdad no prueba atribución'}
            hist.append(row)
    matrix.extend(hist)
    for b in balances:
        relevant=[r for r in matrix if r['anio']==b['anio'] and r['codigo_sies']==b['codigo_sies']]
        b['matriculas_enlazadas']=sum(r.get('matriculas_mu_incluidas',0) for r in relevant if not r.get('sin_enlace_motivo'))
        b['matriculas_sin_enlace']=sum(r.get('matriculas_mu_incluidas',0) for r in relevant if r.get('sin_enlace_motivo'))
        b['sin_enlace_por_motivo']=dict(sorted((r['sin_enlace_motivo'],r['matriculas_mu_incluidas']) for r in relevant if r.get('sin_enlace_motivo') and r['matriculas_mu_incluidas']))
    # Validación de 95 filas: usar asociaciones, no todos los alias de la fila como aprobados.
    validation=[]
    for r in table:
        if r['anio']!=2027 or r['estado'] not in {'CONTINUA','VERSION_MULTIPLE_RESUELTA','APROBADA_D_2027'}: continue
        c=r['codigo_unico']; years=[y for (y,cc),b in bindex.items() if cc==c and b['ancla_codigo']]
        y=max(years) if years else None; expected=set(); historical_groups=set()
        for g in r['candidatos_internos']:
            if g.get('estado_paridad_2027')=='PLAN_HISTORICO_SIN_ALUMNOS_VIGENTES':
                historical_groups.add((g['codcarpr'],g['sede_umas'],g['jornada_umas']));continue
            approved=g.get('codigo_sies_aprobado_grupo')
            sel=g.get('seleccion_grupo',{}).get('seleccion')
            if approved and approved!=c:continue
            if not approved and sel and sel!=c:continue
            for p in g.get('plan_umas',[]): expected.add((g['codcarpr'],p,g['sede_umas'],g['jornada_umas']))
        observed=[x for x in matrix if x['anio']==y and x['codigo_sies']==c and x.get('matriculas_mu_incluidas',0)>0]
        good=[];bad=[];unknown=0
        for x in observed:
            key=(x['codcarpr'],x['plan_umas'],x['sede_umas'],x['jornada_umas'])
            if x['codcarpr']=='SIN_ENLACE':unknown+=x['matriculas_mu_incluidas']
            elif key in expected:good.append(x)
            else:bad.append(x)
        state='CONTRADICE' if bad else ('CONFIRMA' if good and not unknown else 'SIN_DATO_ANUAL')
        validation.append({'codigo_sies':c,'estado_H002':r['estado'],'anio_ancla_mas_reciente':y,'estado':state,
            'coincidentes':sum(x['matriculas_mu_incluidas'] for x in good),'contradictorias':sum(x['matriculas_mu_incluidas'] for x in bad),
            'sin_enlace':unknown,'esperados':sorted(expected),'observados':[{'plan':x['plan_umas'],'codcarpr':x['codcarpr'],
            'sede':x['sede_umas'],'jornada':x['jornada_umas'],'matriculas':x['matriculas_mu_incluidas']} for x in good+bad],
            'clase':'C/E','nota':'Contraste temporal del catálogo, no revocación D ni prueba de que todo alias sea equivalente. Parcial sin enlace no confirma toda la fila.'})
    assert len(validation)==95
    # Resumen por carrera SIES y secuencia anual de códigos, con primeros años observados.
    summary=[]
    for car in sorted(first_sies):
        codes={c for y,c in public if career(c)==car}; cg=set().union(*(fam_groups[family(c)] for c in codes))
        internal={x[0] for x in cg}; first_internal={x:min(first_umas[x]) if first_umas[x] else None for x in sorted(internal)}
        seq=[]
        for y in sorted({y for y,c in public if c in codes}):
            entries=[{'codigo':c,'sies_total':public[(y,c)]['total'],
                      'estados':sorted({x['estado'] for x in matrix if x['anio']==y and x['codigo_sies']==c}),
                      'control_codigo':bindex.get((y,c),{}).get('estado','COMPARACION_AGREGADA_C_E')}
                     for c in sorted(codes) if (y,c) in public]
            seq.append({'anio':y,'codigos':entries})
        summary.append({'cod_carrera_sies':car,'primer_anio_sies_observado':min(first_sies[car]),
                        'primer_anio_umas_observado':min((v for v in first_internal.values() if v is not None),default=None),
                        'primer_anio_umas_por_codcarpr':first_internal,'secuencia':seq,
                        'clase':'B/C/E','limite':'Inicio observado en fuentes disponibles; no inicio oficial. No imputar años entre ingreso y última matrícula.'})
    yearly=[]
    for y in [2025,2024,2023,2022,2026]:
        mm=[x for x in matrix if x['anio']==y]; bb=[b for b in balances if b['anio']==y and b['estado']!='FUERA_ALCANCE_POSTITULO']
        linked=sum(x['matriculas_mu_incluidas'] for x in mm if not x['sin_enlace_motivo']); included=sum(x['matriculas_mu_incluidas'] for x in mm)
        group_states=Counter({})
        keys={(x['codcarpr'],x['plan_umas'],x['sede_umas'],x['jornada_umas'],x['estado']) for x in mm if x['matriculas_mu_incluidas'] and x['codcarpr']!='SIN_ENLACE'}
        group_states.update(k[-1] for k in keys)
        yearly.append({'anio':y,'filas_mu':sum(x['matriculas_mu_incluidas']+x['matriculas_mu_excluidas_vig'] for x in mm),
                       'mu_seleccionada':included,'enlazadas':linked,'porcentaje_enlazado':round(100*linked/included,3),
                       'sin_enlace_matriculas':included-linked,'sin_enlace_codigo_motivo':sum(1 for x in mm if x['sin_enlace_motivo'] and x['matriculas_mu_incluidas']),
                       'combinaciones_plan_conocido':dict(sorted(group_states.items())),
                       'personas_comparables':len(selected_people[y]-{None}),'personas_enlazadas':len(linked_people[y]-{None}),
                       'codigos_con_diferencia':sum(b['estado']=='DIFERENCIA_E' for b in bb),
                       'sies_total_comparable':sum(b['sies_total'] for b in bb),
                       'todos_codigos_cuadran':all(b['estado']=='CUADRA' for b in bb),
                       'ancla':'SI_POR_CODIGO' if y!=2026 and all(b['estado']=='CUADRA' for b in bb) else ('PARCIAL_POR_CODIGO' if y!=2026 else 'NO_2026')})
    admp=[x for x in matrix if x['anio'] in [2025,2026] and family(x['codigo_sies'])=='I162S2C83J4']
    admp_gaps=Counter()
    for d in details:
        if d['anio'] in [2025,2026] and family(d['codigo_sies'])=='I162S2C83J4' and not d['seleccion']:
            cars=','.join(sorted({ca for cl,ca in d['expedientes_DA_familia']})) or 'SIN_CARRERA_DA'
            admp_gaps[(d['anio'],d['codigo_sies'],d['vig'],cars,d['motivo'])]+=1
    priority={'matriz':admp,'controles':[b for b in balances if b['anio'] in [2025,2026] and family(b['codigo_sies'])=='I162S2C83J4'],
              'sin_enlace_desglose':[{'anio':y,'codigo':c,'vig':v,'carreras_DA':ca,'motivo':why,'matriculas':n}
                                    for (y,c,v,ca,why),n in sorted(admp_gaps.items())],
              'dictamen':'ADMP→V1: respaldo anual parcial en 2025. CADMP→V2: respaldo anual parcial en 2026, que no es año ancla institucional. Los enlaces faltantes impiden cerrar íntegramente los dos códigos; no se fuerza el plan desde el total. Se conserva H002 y su pendiente E. V2 no tuvo matrícula observada en 2025 (ausencia, no fila cero inventada).'}
    report={'clases':'A norma; B observado; C cálculo; D decisión; E pendiente','anual':yearly,'ADMP_CADMP':priority,
            'validacion_95':dict(Counter(x['estado'] for x in validation)),
            'historia_estados_filas':dict(Counter(x['estado'] for x in hist)),
            'historia_carreras_por_estado_no_sumar':{s:len({career(x['codigo_sies']) for x in hist if x['estado']==s}) for s in ['CUADRA_TOTAL','APROXIMA','NO_CUADRA']},
            'cobertura_fuentes':coverage,'extractos_DA_no_utilizados':skipped,'actividad_diagnostico':dict(activity_diagnostics),
            'periodos_literales':dict(periods),'codigos_publicos_no_ISCJV':nonstandard,'puente_reproducido':bridge_ok,
            'pendiente_cierre_historico':'No se declara completa la atribución: actividad anual y catálogo no certifican matrícula al corte, ni períodos sin cobertura.'}
    outputs={DATA/'MATRIZ_PARIDAD_ANUAL.json':{'filas':matrix},DATA/'RESUMEN_POR_CARRERA.json':{'carreras':summary},
             OUT/'CONTROL_POR_CODIGO_ANIO.json':{'filas':balances},OUT/'VALIDACION_95_H002.json':{'filas':validation},OUT/'RESULTADOS.json':report}
    return sources, outputs


def csv_text(rows):
    keys=list(dict.fromkeys(k for r in rows for k in r)); f=io.StringIO(newline='');w=csv.DictWriter(f,fieldnames=keys,delimiter=';',lineterminator='\n');w.writeheader()
    for r in rows:
        w.writerow({k:json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(dict,list,tuple)) else v for k,v in r.items()})
    return f.getvalue()


def render(outputs):
    report=outputs[OUT/'RESULTADOS.json'];matrix=outputs[DATA/'MATRIZ_PARIDAD_ANUAL.json']['filas'];validation=outputs[OUT/'VALIDACION_95_H002.json']['filas']
    lines=['# H002f · paridad anual (propuesta, no adjudicación)','',
           '## ADMP / CADMP · prioridad','',
           '| Año | Código | CODCARPR / plan | MU incluidas | MU excluidas VIG | Estado |',
           '|---|---|---|---:|---:|---|']
    for x in report['ADMP_CADMP']['matriz']:
        lines.append(f"| {x['anio']} | {x['codigo_sies']} | {x['codcarpr']} / {x['plan_umas']} | {x['matriculas_mu_incluidas']} | {x['matriculas_mu_excluidas_vig']} | {x['estado']} |")
    lines += ['', '| Año | Código | MU seleccionada | SIES | Diferencia |','|---|---|---:|---:|---:|']
    for x in report['ADMP_CADMP']['controles']:
        lines.append(f"| {x['anio']} | {x['codigo_sies']} | {x['mu_seleccionada']} | {x['sies_total']} | {x['diferencia_sies_menos_mu']} |")
    lines += ['',report['ADMP_CADMP']['dictamen'],'','Sin enlace de plan:','']
    for x in report['ADMP_CADMP']['sin_enlace_desglose']:
        lines.append(f"- {x['anio']} · {x['codigo']} · VIG {x['vig']}: {x['matriculas']} matrículas; carreras DA {x['carreras_DA']}; {x['motivo']}.")
    lines += ['','## Controles por año','',
              '| Año | MU filas | Selección | Enlazadas | % enlace | Sin enlace | Códigos distintos | Ancla |',
              '|---|---:|---:|---:|---:|---:|---:|---|']
    for x in report['anual']:
        lines.append(f"| {x['anio']} | {x['filas_mu']} | {x['mu_seleccionada']} | {x['enlazadas']} | {x['porcentaje_enlazado']} | {x['sin_enlace_matriculas']} | {x['codigos_con_diferencia']} | {x['ancla']} |")
    lines += ['C: % enlace sobre matrículas seleccionadas, no sobre personas. SIN_ENLACE no identifica un plan conocido: se cuenta por código y motivo, separado de combinaciones de plan. Personas no aditivas.', '',
              '## Clasificación anual de combinaciones conocidas','']
    for x in report['anual']:lines.append(f"- {x['anio']}: {json.dumps(x['combinaciones_plan_conocido'],ensure_ascii=False)}; SIN_ENLACE: {x['sin_enlace_codigo_motivo']} combinaciones código × motivo, {x['sin_enlace_matriculas']} matrículas.")
    lines += ['', '## Validación de 95 códigos H002', '', json.dumps(report['validacion_95'],ensure_ascii=False), '',
              'CONFIRMA exige todos los enlaces del código en el ancla sin contradicción; SIN_DATO_ANUAL incluye cobertura parcial. CONTRADICE es diferencia temporal con la tabla de 2027 para revisión, no revoca aprobaciones ni acredita error.', '',
              '| Código H002 | Último ancla del código | Resultado | Coincidentes | Contradictorias | Sin enlace |',
              '|---|---|---|---:|---:|---:|']
    for v in validation:
        lines.append(f"| {v['codigo_sies']} | {v['anio_ancla_mas_reciente']} | {v['estado']} | {v['coincidentes']} | {v['contradictorias']} | {v['sin_enlace']} |")
    lines += ['', '## Antes de 2022', '',json.dumps(report['historia_carreras_por_estado_no_sumar'],ensure_ascii=False), '',
              'Carreras por estado no sumables: una carrera puede tener varios años y planes con estados diferentes. Primeros años y secuencia completa en RESUMEN_POR_CARRERA.csv. C/E: igualdad entre actividad y matrícula no certifica identidad ni versión.', '',
              '| Carrera SIES | Primer año SIES observado | Primer año UMAS observado | Estados anteriores a 2022 |',
              '|---|---:|---:|---|']
    for s in outputs[DATA/'RESUMEN_POR_CARRERA.json']['carreras']:
        states=sorted({x['estado'] for x in matrix if x['anio']<2022 and career(x['codigo_sies'])==s['cod_carrera_sies']})
        lines.append(f"| {s['cod_carrera_sies']} | {s['primer_anio_sies_observado']} | {s['primer_anio_umas_observado']} | {', '.join(states) or 'SIN HISTORIA PREVIA EN FUENTE'} |")
    lines += ['',
              '## Sumas por código y año', '', '| Año | Código | MU seleccionada | SIES | Diferencia | Estado |','|---|---|---:|---:|---:|---|']
    for b in outputs[OUT/'CONTROL_POR_CODIGO_ANIO.json']['filas']:
        lines.append(f"| {b['anio']} | {b['codigo_sies']} | {b['mu_seleccionada']} | {b['sies_total']} | {b['diferencia_sies_menos_mu']} | {b['estado']} |")
    lines += ['', 'Los filtros VIG son el patrón B/C del puente, no una norma universal. 2025 incorpora las 73 VIG 0 acreditadas en precarga AC2026; 2026 es contraste y conserva el cierre documentado de la brecha.', '',
              '## Fuentes y límites', '', 'FUENTES.json fija cada ruta y SHA-256; MANIFEST registra generador y resultados. Todas las lecturas son locales, sin modificar originales.', '',
              report['pendiente_cierre_historico'], '', 'Datos Alumnos aporta identidad/expediente y extremos de año; sus fotos no certifican sede o estado al corte. PROMEDIOS acredita actividad anual en los períodos originales. La unión de extractos evita doble conteo, no resuelve modificaciones históricas. No se asigna una versión por similitud de totales.', '']
    # Lotes estables por carrera SIES, con ancla más reciente primero.
    issues=defaultdict(list)
    for x in matrix:
        if x['estado'] in {'REPARTIDO','SIN_ENLACE','UNICO_SIN_CUADRATURA','NO_CUADRA','APROXIMA'} and (x.get('matriculas_mu_incluidas',1)>0):
            issues[career(x['codigo_sies'])].append(x)
    for v in validation:
        if v['estado']=='CONTRADICE':issues[career(v['codigo_sies'])].append({'anio':v['anio_ancla_mas_reciente'],'codigo_sies':v['codigo_sies'],'estado':'CONTRADICE_H002','detalle':v})
    ll=['# Lotes de revisión · cinco carreras SIES por lote','','C/E: propuestas no aplicadas. Primero el ancla más reciente, después contraste 2026 e historia descendente.','']
    cars=sorted(issues)
    for start in range(0,len(cars),5):
        ll += [f'## Lote {start//5+1}','']
        for car in cars[start:start+5]:
            ll += [f'### Carrera {car}','']
            rows=sorted(issues[car],key=lambda x:(not x.get('ancla_codigo',False),-int(x['anio']),x['codigo_sies'],x.get('plan_umas','')))
            for x in rows:
                n=x.get('matriculas_mu_incluidas',x.get('expedientes_actividad_anual','ver detalle'))
                ll += [f"- {x['anio']} · {x['codigo_sies']} · {x.get('codcarpr','')} / {x.get('plan_umas','')} · {x['estado']} · conteo {n}; SIES {x.get('sies_codigo_total','ver VALIDACION_95_H002.json')}. {x.get('sin_enlace_motivo','')}"]
                if x['estado']=='CONTRADICE_H002':
                    for z in x['detalle']['observados']:
                        ll.append(f"  Observado: {z['codcarpr']} / {z['plan']}, {z['sede']}, {z['jornada']}: {z['matriculas']} matrículas.")
                ll.append('  Decisión Alexi: ____')
            ll.append('')
    report['lotes_revision']={'carreras':len(cars),'lotes':(len(cars)+4)//5,'casos':sum(map(len,issues.values()))}
    return '\n'.join(x.rstrip() for x in lines).rstrip()+'\n', '\n'.join(x.rstrip() for x in ll).rstrip()+'\n'


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--verificar',action='store_true');args=parser.parse_args()
    sources,outputs=build();report,lotes=render(outputs)
    texts={p:js(v) for p,v in outputs.items()}
    texts[OUT/'INFORME.md']=report;texts[OUT/'LOTES_REVISION.md']=lotes
    for name,key in [('MATRIZ_PARIDAD_ANUAL','filas'),('RESUMEN_POR_CARRERA','carreras')]:texts[DATA/(name+'.csv')]=csv_text(outputs[DATA/(name+'.json')][key])
    texts[DATA/'README.md']=(OUT/'ESPECIFICACION.md').read_text()+'''\n## Reproducción\n\nEjecutar `python3 outputs/paridad_anual_20261029/construir_anual.py --verificar` desde el repositorio con las fuentes locales de FUENTES.json. Verifica hashes y reproduce todas las salidas; las identidades solo se escriben en el directorio privado.\n\n## Estados y unidades\n\nCUADRA_UNICO / REPARTIDO / SIN_ENLACE: enlace anual MU; UNICO_SIN_CUADRATURA evita confirmar código con total diferente. FUERA_SELECCION_VIG conserva filas no incluidas. Antes de 2022: CUADRA_TOTAL / APROXIMA (±10 %) / NO_CUADRA son contrastes C/E de expedientes con actividad versus matrículas, no asignación. Las filas de comparación histórica repiten el total público por candidato y NO son sumables. Los conteos MU incluidos/excluidos sí son aditivos por código y año. Los totales SIES repetidos en la matriz se consultan una vez en CONTROL_POR_CODIGO_ANIO.json.\n\n## Los 37 códigos sin historia\n\nD · Alexi, 29-10-2026: se revisan con este motor cuando exista MU 2027 (o MU del año de su primer ingreso), por código y total y con enlace de plan inequívoco. No se solicitan a UMAS. Hasta entonces conservan NUEVO_2027_SIN_PLAN_E. Un total coincidente sin enlace no cierra la identidad del plan. No se modifica H002 en este encargo.\n'''
    manifest={'fecha':'2026-10-29','clase':'C','base_git':'507231b9128db76c17a470c91b18869a4adff425','fuentes':sources,
              'generador':{'ruta':str(Path(__file__).relative_to(ROOT)),'sha256':sha(__file__)},
              'auxiliares':{str(p.relative_to(ROOT)):sha(p) for p in [OUT/'FUENTES.json',OUT/'ESPECIFICACION.md',OUT/'ESTADO_INICIAL.json',OUT/'test_anual.py',OUT/'qa_anual.py']},
              'salidas':{str(p.relative_to(ROOT)):hashlib.sha256(t.encode()).hexdigest() for p,t in sorted(texts.items())}}
    texts[DATA/'MANIFEST.json']=js(manifest)
    for p,t in texts.items():
        if args.verificar:
            if not p.exists() or p.read_bytes()!=t.encode():raise ValueError('Reproducción distinta: '+str(p.relative_to(ROOT)))
        else:
            p.parent.mkdir(parents=True,exist_ok=True);p.write_text(t)
    print(js({'reproduccion':'PASS' if args.verificar else 'GENERADO','anual':outputs[OUT/'RESULTADOS.json']['anual'],'validacion_95':outputs[OUT/'RESULTADOS.json']['validacion_95']}))


if __name__=='__main__':main()
