"""Audita identidad histórica H113 en memoria, sin modificar fuentes ni motores."""
from pathlib import Path
import re,json,hashlib,subprocess,sys,sqlite3
from collections import Counter,defaultdict
from datetime import datetime,timezone
import pandas as pd
import openpyxl
ROOT=Path(__file__).resolve().parents[2]; OUT=Path(__file__).parent
OLD=ROOT/'outputs/avance_historico_origen_h113_20261027'
manifest=json.loads((OLD/'MANIFEST_ORIGEN_H113.json').read_text())
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(n,x):(OUT/n).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def git(*a):return subprocess.check_output(['git',*a],cwd=ROOT).decode().strip()
def body(x):return re.sub(r'[.\s]','',str(x)).upper().strip()
def legacy(x):
 s=re.sub('[^0-9K]','',str(x).upper().replace('.','').replace('-','').strip())
 return s[:-1] if len(s)>1 else s
src={k:Path(v['ruta']) for k,v in manifest['fuentes'].items()}
checks=[]
def ck(n,e,v):checks.append({'control':n,'esperado':e,'observado':v,'estado':'PASS' if e==v else 'FAIL','clase':'B/C'})
start={'head':git('rev-parse','HEAD'),'branch':git('branch','--show-current')}
# Freeze existing modified tracked files; do not change them.
protected={x:sha(ROOT/x) for x in git('diff','--name-only','HEAD').splitlines() if (ROOT/x).is_file()}
for k,p in src.items():ck('Hash entrada '+k,manifest['fuentes'][k]['sha256'],sha(p))
cols=['TIPO_DOCUMENTO','NUM_DOCUMENTO','DV','PRIMER_APELLIDO','SEGUNDO_APELLIDO','NOMBRES','SEXO','FECHA_NACIMIENTO','CODIGO_UNICO','PLAN_ESTUDIOS','ANIO_INGRESO_CARRERA_ACTUAL','SEM_INGRESO_CARRERA_ACTUAL','ANIO_INGRESO_CARRERA_ORIGEN','SEM_INGRESO_CARRERA_ORIGEN','CURSO_1ER_SEM','CURSO_2DO_SEM','UNIDADES_CURSADAS','UNIDADES_APROBADAS','UNID_CURSADAS_TOTAL','UNID_APROBADAS_TOTAL','VIGENCIA']
measures=cols[14:20]
target=pd.read_csv(src['h113'],sep=';',header=None,names=cols,dtype=str,keep_default_na=False,encoding='latin1')
frames=pd.read_excel(src['promedios'],sheet_name=None,dtype=str,keep_default_na=False)
print('Fuentes leidas; identidades solo en memoria',flush=True)
parts=[]
for name,f in frames.items():
 f=f.copy();f['_HOJA_PROMEDIOS']=name;parts.append(f)
prom=pd.concat(parts,ignore_index=True,sort=False)
prom['_CODCLI_NORM']=prom['CODCLI'].astype(str).str.strip()
prom['_DOC_NORM']=prom['RUT'].map(legacy)
prom['_ANO_NUM']=pd.to_numeric(prom['ANO'].astype(str).str.extract(r'(\d{4})',expand=False),errors='coerce')
prom['_PERIODO_NUM']=pd.to_numeric(prom['PERIODO'].astype(str).str.replace(',','.'),errors='coerce')
prom['_ESTADO_NORM']=prom['ESTADO'].fillna('').str.upper().str.strip()
prom=prom.drop_duplicates().reset_index(drop=True)
# Diagnostic keys added after full-row de-duplication, as in the earlier replay.
prom['_BODY']=prom['RUT'].fillna('').map(body)
idxlegacy={k:set(v) for k,v in prom.groupby('_DOC_NORM').groups.items()}
idxbody={k:set(v) for k,v in prom.groupby('_BODY').groups.items()}
idxcli={k:set(v) for k,v in prom.groupby('_CODCLI_NORM').groups.items()}
def select(key,index):
 ix=set(index.get(key,set())) if key else set()
 for cli in set(prom.loc[list(ix),'_CODCLI_NORM']):
  if cli and cli!='nan':ix|=idxcli.get(cli,set())
 return sorted(ix)
def calc(ix):
 s=prom.loc[ix];ann=s[s['_ANO_NUM'].eq(2025)];hist=s[s['_ANO_NUM'].le(2025)]
 ar=ann[ann['_ESTADO_NORM'].isin(['A','R'])];a=ann[ann['_ESTADO_NORM'].eq('A')]
 per=set(ar['_PERIODO_NUM'].dropna())
 return [int(bool(per & {1,5})),int(bool(per & {2,3,4,6})),len(ar),len(a),int(hist['_ESTADO_NORM'].isin(['A','E','I','R']).sum()),int(hist['_ESTADO_NORM'].isin(['A','E','I']).sum())]
def expected(r):return [int(r[c]=='SI') if i<2 else int(r[c]) for i,c in enumerate(measures)]
records=target.to_dict('records'); oldvals=[];newvals=[];missing=[]
for r in records:
 ix=select(legacy(r['NUM_DOCUMENTO']),idxlegacy);jx=select(body(r['NUM_DOCUMENTO']),idxbody)
 a,b=calc(ix),calc(jx);oldvals.append(a);newvals.append(b)
 if a!=expected(r):missing.append((r,ix,jx))
ck('Reproduccion coincide auditoria anterior',2370,sum(a==expected(r) for a,r in zip(oldvals,records)))
bygroup=defaultdict(list)
for i,r in enumerate(records):bygroup[legacy(r['NUM_DOCUMENTO'])].append(i)
collisiongroups=[v for v in bygroup.values() if len({(records[i]['TIPO_DOCUMENTO'],records[i]['NUM_DOCUMENTO'],records[i]['DV']) for i in v})>1]
collisions=[]
for j,indices in enumerate(collisiongroups,1):
 collisions.append({'grupo':f'COL-{j:02d}','filas_precarga':len(indices),'codigos_sies':sorted({records[i]['CODIGO_UNICO'] for i in indices}),'filas_con_cambio_en_conteos':sum(oldvals[i]!=newvals[i] for i in indices),'unidades_antes':[sum(oldvals[i][k] for i in indices) for k in range(2,6)],'unidades_escenario_cuerpo':[sum(newvals[i][k] for i in indices) for k in range(2,6)],'clase':'B/C/E'})
changed=[i for i in range(len(records)) if oldvals[i]!=newvals[i]]
# Strict body+DV diagnostic solely on actual subject source (Hoja1).
h=frames['Hoja1'];strict_index=defaultdict(list)
for i,r in h.iterrows():strict_index[(body(r['RUT']),body(r['DIG']))].append(i)
strict_coverage=Counter(); alerts=[]
for i in changed:
 r=records[i]; ids=strict_index.get((body(r['NUM_DOCUMENTO']),body(r['DV'])),[])
 strict_coverage['con_asignaturas_cuerpo_y_dv_exactos' if ids else 'sin_asignaturas_cuerpo_y_dv_exactos']+=1
# Read common dictionary with its own hash checks; do not invent new identity links.
sys.path.insert(0,str(ROOT))
from common.diccionario_sies_seguimiento import DiccionarioSiesSeguimiento
service=DiccionarioSiesSeguimiento()
missingresults=[]
for r,ix,jx in missing:
 d={'clase':'B/C/E','tipo_documento':r['TIPO_DOCUMENTO'],'codigo_sies':r['CODIGO_UNICO'],'plan_sies':r['PLAN_ESTUDIOS'],'dv_presente':bool(r['DV']),'coincidencias_cuerpo_por_hoja':{name:int(f['RUT'].map(body).eq(body(r['NUM_DOCUMENTO'])).sum()) for name,f in frames.items() if 'RUT'in f},'h113':dict(zip(measures,expected(r))),'legacy':dict(zip(measures,calc(ix))),'cuerpo':dict(zip(measures,calc(jx)))}
 if r['DV']:
  try:
   val=service.codigos_mu(persona=body(r['NUM_DOCUMENTO'])+'-'+body(r['DV']),anio=2025)
   d['codigos_mu_2025']=val['codigos']
   persona=body(r['NUM_DOCUMENTO'])+'-'+body(r['DV'])
   rows=service.db.execute('SELECT carrera,plan,anio,periodo,payload FROM observaciones WHERE persona=?',(persona,)).fetchall()
   d['contextos_diccionario']=len(rows)
   d['planes_diccionario']=sorted({v[1] for v in rows});d['carreras_diccionario']=sorted({v[0] for v in rows})
  except ValueError:d['consulta_diccionario']='Identidad no compatible con contrato RUT+DV; no se fuerza'
 else:d['consulta_diccionario']='DV ausente; contrato no permite inventarlo'
 missingresults.append(d)
service.close()
# Explicit correspondence in governed MU processing workbook, never fuzzy names.
bridge_path=src['promedios'].parents[1]/'02_MATRICULA_UNIFICADA_Y_CARGAS_SIES/IPSS_Procesamiento.xlsx'
bridge_sha=sha(bridge_path)
ck('Hash puente catalogo gobernado','e2984ccdf22ecf35c46e1f77275fdd12cda47bd2d77a523860e2cd2cf90127f4',bridge_sha)
wb=openpyxl.load_workbook(bridge_path,read_only=True,data_only=True)
ws=wb['2025_carga']; br=list(ws.iter_rows(values_only=True)); finalrows=list(wb['Final a carga'].iter_rows(values_only=True)); wb.close()
header=br[0]
ck('Puente encabezados J BU BV', ['RUT','N_DOC','DV'],[header[9],header[72],header[73]])
def cellstr(v):
 if v is None:return ''
 if isinstance(v,float) and v.is_integer():return str(int(v))
 return body(v)
bridge_info=[]; replacements={}
for missing_index,(r,ix,jx) in enumerate(missing):
 matches=[(lineno,row) for lineno,row in enumerate(br[1:],2) if cellstr(row[72])==body(r['NUM_DOCUMENTO'])]
 fhead=finalrows[0]; fi_doc=fhead.index('N_DOC'); fi_type=fhead.index('TIPO_DOC'); fi_dv=fhead.index('DV')
 finalmatches=[f for f in finalrows[1:] if cellstr(f[fi_doc])==body(r['NUM_DOCUMENTO'])]
 internal={cellstr(row[9]) for _,row in matches if row[9] is not None}
 codes=sorted({'I'+cellstr(row[81])+'S'+cellstr(row[82])+'C'+cellstr(row[83])+'J'+cellstr(row[85])+'V'+cellstr(row[86]) for _,row in matches})
 b={'clase':'B/C/E','hoja':'2025_carga','columnas':['J RUT','BU N_DOC','BV DV','CD:CI componentes SIES'],'filas_coincidentes':len(matches),'identificadores_internos_distintos':len(internal),'codigos_observados':codes,'codigo_h113_presente':r['CODIGO_UNICO'] in codes,'resultado':'PENDIENTE','tipo_h113':r['TIPO_DOCUMENTO'],'tipos_final_carga':sorted({cellstr(f[fi_type]) for f in finalmatches}),'dv_h113_presente':bool(r['DV']),'dv_puente_presente':all(bool(cellstr(row[73])) for _,row in matches),'dv_igual_h113':all(cellstr(row[73])==body(r['DV']) for _,row in matches),'criterio':'N_DOC exacto + codigo completo; la discrepancia tipo/DV no se resuelve ni se omite'}
 if len(internal)==1 and r['CODIGO_UNICO'] in codes:
  full=next(iter(internal));parts=full.rsplit('-',1)
  ibody=parts[0];idv=parts[1] if len(parts)==2 else None
  ids=select(ibody,idxbody)
  hmatch=h[h['RUT'].map(body).eq(ibody)]
  dvs={body(x) for x in hmatch['DIG']}
  b.update({'rut_interno_con_dv_explicito':idv is not None,'dv_coincide_promedios':idv in dvs if idv else None,'codcli_en_asignaturas':int(hmatch['CODCLI'].nunique()),'planes_en_asignaturas':sorted(set(hmatch['PLAN_DE_ESTUDIO'])),'carreras_en_asignaturas':sorted(set(hmatch['CODCARR'])),'valores_por_enlace':dict(zip(measures,calc(ids))),'reproduce_seis_columnas':calc(ids)==expected(r)})
  if calc(ids)==expected(r):
   b['resultado']='ESCENARIO_NUMERICO_REPRODUCE_CON_DISCREPANCIA_TIPO_DV_PENDIENTE'
   replacements[next(i for i,x in enumerate(records) if x==r)]=calc(ids)
 bridge_info.append(b)
legacy_bridge=[replacements.get(i,v) for i,v in enumerate(oldvals)]
body_bridge=[replacements.get(i,v) for i,v in enumerate(newvals)]
ck('Puente fuente intacta',bridge_sha,sha(bridge_path))
result={'clase':'B/C/E','universo':len(target),'denominadores_modificados':False,'escenario_adoptado':False,'campos':measures,'conciliacion_legacy':{'coinciden':2370,'difieren':len(missing)},'escenario_cuerpo_completo':{'registros_con_cambio_vs_legacy':len(changed),'sumas_h113':dict(zip(measures[2:],[sum(expected(r)[k] for r in records) for k in range(2,6)])),'sumas_legacy':dict(zip(measures[2:],[sum(v[k] for v in oldvals) for k in range(2,6)])),'sumas_cuerpo':dict(zip(measures[2:],[sum(v[k] for v in newvals) for k in range(2,6)])),'cobertura_dv_en_filas_que_cambian':dict(strict_coverage)},'siete_grupos':collisions,'registro_diferente':missingresults,'limites':['Escenario de identidad exclusivamente; conserva seleccion multicarrera del historico.','Cuerpo sin DV no es cierre de identidad; DV se contrasta separadamente.','No se recupero MAPEO_IDENTIDAD_CODCLI.xlsx; no se afirma reproduccion integral.','No hay datos individuales en las salidas.']}
result['puente_explicito_mu']=bridge_info
result['escenario_condicionado_legacy_mas_puente']={'coinciden_seis':sum(v==expected(r) for v,r in zip(legacy_bridge,records)),'sumas':dict(zip(measures[2:],[sum(v[k] for v in legacy_bridge) for k in range(2,6)]))}
result['escenario_condicionado_cuerpo_mas_puente']={'sumas':dict(zip(measures[2:],[sum(v[k] for v in body_bridge) for k in range(2,6)])),'filas_con_cambio_vs_h113':sum(v!=expected(r) for v,r in zip(body_bridge,records))}
dump('RESULTADOS.json',result)
for k,p in src.items():ck('Fuente intacta '+k,manifest['fuentes'][k]['sha256'],sha(p))
for f,x in protected.items():ck('Cambio ajeno intacto '+f,x,sha(ROOT/f))
ck('HEAD intacto',start['head'],git('rev-parse','HEAD'));ck('Rama intacta',start['branch'],git('branch','--show-current'))
dump('QA.json',{'estado':'PASS_CONTROLES_ESCENARIO_NO_ADOPTADO' if all(v['estado']=='PASS' for v in checks) else 'FAIL','controles':checks,'pendiente_atribucion_academica':True})
dump('MANIFEST.json',{'fecha_utc':datetime.now(timezone.utc).isoformat(),'fuentes':{**manifest['fuentes'],'puente_mu':{'ruta':str(bridge_path),'sha256':bridge_sha}},'script_sha256':sha(__file__),'git_antes':start,'git_despues':{'head':git('rev-parse','HEAD'),'branch':git('branch','--show-current')},'fuente_auditoria_previa':{'ruta':str(OLD/'MANIFEST_ORIGEN_H113.json'),'sha256':sha(OLD/'MANIFEST_ORIGEN_H113.json')},'diccionario_descriptor':{'edicion':service.edicion,'base_sha256':service.descriptor['base_privada']['sha256']},'salidas':{n:sha(OUT/n) for n in ['RESULTADOS.json','QA.json']},'commit':False,'push':False})
print(json.dumps(result,ensure_ascii=False,indent=2))
