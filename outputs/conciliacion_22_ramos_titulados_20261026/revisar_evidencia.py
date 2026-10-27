"""Contraste agregado de ramos y titulados sin exportar identificadores."""
import collections,csv,hashlib,json,re,sqlite3,datetime
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
prior=ROOT/'outputs/conciliacion_mallas_vigencia_20261026/REVISION_AGREGADA.json'
structure=ROOT/'data/versiones_historicas_sies_umas/fase_4_20261015/estructura_asignaturas.json'
audit=ROOT/'data/versiones_historicas_sies_umas/fase_4_20261015/auditoria_catalogo.json'
ones=Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Datos DAI (RAW) - 1.1 Datos crudos/1.2.1 Datos Institucionales/1.2.1.1 SIES/Enviados')
book=next(ones.glob('Matri* Unificada/2026/SEGUNDA CARGA/ARCHIVOS VARIOS/1668382*xlsx'))
titlecsv=next(ones.glob('Titulaci*/2026/Titulaci* Unificada 2026.csv'))
db=Path('/Users/alexi/.local/share/ipss/v11_fase2_20261026/DICCIONARIO_SIES_CANDIDATO_V11_86_Y_AMBIGUEDAD.sqlite')
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for v in iter(lambda:f.read(1048576),b''):h.update(v)
 return h.hexdigest()
assert sha(book)=='149367ca3f7be7717f958e2ce0f8924e18353e6302766b1abf737c54b5dce8e3'
assert sha(titlecsv)=='b52586f6741109e293c90896721ca5b879f5f51ec98a1c80650464aa52356dec'
assert sha(db)=='1dafb1486f3ebd080ad4435a3af80b65a1a7016ba47b51f30d91b9a050985141'
def s(x):
 if x is None:return ''
 if isinstance(x,(float,int)) and x==int(x):return str(int(x))
 return str(x).strip().upper()
def rut(n,d):
 a=s(n).replace('.','').replace(' ','');b=s(d)
 return str(int(a))+'-'+b if a.isdigit() and re.fullmatch('[0-9K]',b) else None
p=json.loads(prior.read_text());struct=json.loads(structure.read_text());a=json.loads(audit.read_text());links=collections.defaultdict(list)
for row in struct:
 for e in row.get('equivalencias_detalle',[]):
  links[(s(row['CODPESTUD']),s(row['CODRAMO']))].append({'plan':s(e['plan_destino']),'ramo':s(e['ramo_destino']),'filas_fuente':e['filas_fuente'],'literal':e['literal']})
issues=[]
for issue in p['discrepancias']:
 target=(issue['plan'],issue['asignatura']);direct=[]
 for (pl,code),rels in links.items():
  for rel in rels:
   if (pl==target[0] and rel['ramo']==target[1]) or ((pl,code)==target):
    direct.append({'plan_origen':pl,'ramo_origen':code,**rel})
 issues.append({**issue,'relaciones_catalogo':direct,'clasificacion_relacion':'B/C relación explícita en catálogo; E aplicación individual y vigencia' if direct else 'E: sin relación explícita para este contexto'})
w=openpyxl.load_workbook(book,read_only=True,data_only=True)
def rows(sheet):
 it=w[sheet].iter_rows(values_only=True);head=[s(x) for x in next(it)]
 for r in it:yield dict(zip(head,r))
titles=[]
for r in rows('Res_Tit_Carga'):
 ident=rut(r.get('NUM_DOCUMENTO'),r.get('DV'));code=s(r.get('COD_SIES_OBT_TIT'))
 if ident and code and s(r.get('TIPO_DOCUMENTO'))=='R':titles.append((ident,code,s(r.get('FECHA_OBT_TIT_GRA')),s(r.get('TIPO_DOCUMENTO'))))
hist=[]
for r in rows('5319 Titulados Histór'):
 ident=rut(r.get('NUM_DOCUMENTO'),r.get('DV'));code=s(r.get('COD_UNICO_OBT_TIT'));year=s(r.get('AÑO_TITULACION'))
 if ident and code and year and s(r.get('TIPO_DOCUMENTO'))=='R':hist.append((ident,code,year,s(r.get('TIPO_DOCUMENTO'))))
w.close()
# La comparación CSV/XLSX se limita a las posiciones de documento y código,
# verificadas contra Res_Tit_Carga. No afirma igualdad de los 23 campos.
raw=[]
with titlecsv.open(encoding='utf-8-sig',newline='') as f:
 for r in csv.reader(f,delimiter=';'):
  assert len(r)==23
  ident=rut(r[2],r[3]);code=s(r[10]);typ=s(r[1]);
  if ident and code:raw.append((ident,code,typ))
rawkeys=set(raw);bookkeys={(i,c,t) for i,c,_,t in titles}
target_plans={(x['carrera'],x['plan']) for x in issues}
prom_meta=json.loads((ROOT/'outputs/qa_promedios_cobertura_20261026/RESULTADOS_AGREGADOS.json').read_text())['fuentes']['v11_12166']
prom=Path(prom_meta['ruta']);assert sha(prom)==prom_meta['sha256']
pw=openpyxl.load_workbook(prom,read_only=True,data_only=True);it=pw['Hoja1'].iter_rows(values_only=True);head={s(v):i for i,v in enumerate(next(it)) if v is not None};activity=collections.defaultdict(set)
for r in it:
 career,plan,year,period=(s(r[head[k]]) for k in ('CODCARR','PLAN_DE_ESTUDIO','ANO','PERIODO'))
 if (career,plan) in target_plans and year.isdigit() and period.isdigit():activity[(career,plan)].add((int(year),int(period)))
pw.close();assert sha(prom)==prom_meta['sha256']
con=sqlite3.connect(db.as_uri()+'?mode=ro',uri=True);contexts=collections.defaultdict(lambda:collections.defaultdict(set));expedientes=collections.defaultdict(set)
for person,codcli,career,plan,payload in con.execute('select persona,codcli,carrera,plan,payload from observaciones'):
 key=(career,plan)
 o=json.loads(payload);typ=o.get('tipo_doc_normalizado');code=o.get('codigo_documental_atribuible') or o.get('codigo_interno_adoptado')
 if typ=='RUT' and code:
  expedientes[(person,code)].add((codcli,career,plan))
  if key in target_plans:contexts[(person,key)][code].add('DOCUMENTAL' if o.get('codigo_documental_atribuible') else 'INTERNO')
con.close()
perplan=collections.defaultdict(lambda: {'personas_mismo_rut_codigo_titulo':set(),'multiples_expedientes':set(),'codigos':collections.Counter(),'tipos_vinculo':collections.Counter(),'anos_titulo':collections.Counter()})
for ident,code,year,typ in titles+hist:
 for (person,key),codes in contexts.items():
  if person!=ident or code not in codes:continue
  q=perplan[key];q['personas_mismo_rut_codigo_titulo'].add(ident);q['codigos'][code]+=1;q['anos_titulo'][year[:4]]+=1
  for tipo in codes[code]:q['tipos_vinculo'][tipo]+=1
  if len(expedientes[(person,code)])>1:q['multiples_expedientes'].add(ident)
summary=[]
for k in sorted(target_plans):
 q=perplan[k];summary.append({'carrera':k[0],'plan':k[1],'periodos_actividad_promedios':[{'ano':y,'periodo':p} for y,p in sorted(activity[k])],'personas_con_mismo_rut_y_codigo_v11_y_registro_titulacion':len(q['personas_mismo_rut_codigo_titulo']),'personas_con_multiples_expedientes_compatibles':len(q['multiples_expedientes']),'registros_titulo_por_codigo':dict(q['codigos']),'vinculos_v11_por_tipo_no_exclusivos':dict(q['tipos_vinculo']),'fechas_o_anos_titulo_observados':dict(q['anos_titulo']),'limite':'Titulación y código observados: no prueban automáticamente plan efectivo o aceptación. Los registros pueden solaparse entre fuentes; los años no establecen fecha de vigencia del plan.'})
output={'clasificacion':'B registros; C cruce; E vigencia y acreditación','fuentes':{'estructura':{'ruta':str(structure),'sha256':sha(structure),'procedencia':a['fuente']},'respaldo_xlsx':{'ruta':str(book),'sha256':sha(book),'hojas':['5319 Titulados Histór','Res_Tit_Carga']},'csv':{'ruta':str(titlecsv),'sha256':sha(titlecsv)},'promedios':{'ruta':str(prom),'sha256':sha(prom),'hoja':'Hoja1'},'v11_sha256':sha(db)},'discrepancias':issues,'titulados_por_plan':summary,'qa_csv_vs_xlsx':{'csv_filas_con_identidad_codigo':len(raw),'xlsx_filas_con_identidad_codigo':len(titles),'llaves_distintas_csv':len(rawkeys),'llaves_distintas_xlsx':len(bookkeys),'llaves_comunes':len(rawkeys&bookkeys),'csv_sin_coincidencia':len(rawkeys-bookkeys),'xlsx_sin_coincidencia':len(bookkeys-rawkeys),'limite':'Contraste solo por tipo documento, número, DV y código de título; no prueba igualdad íntegra ni aceptación.'},'limitaciones':'Identidad exacta no demuestra plan individual de titulación: se reutiliza vínculo v11 y se mantienen colisiones entre expedientes. No modificar fuentes ni denominadores.'}
(OUT/'EVIDENCIA_AGREGADA.json').write_text(json.dumps(output,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'ramos':len(issues),'relaciones_explicitas':sum(bool(x['relaciones_catalogo']) for x in issues),'planes':len(summary),'planes_con_coincidencias_titulo':sum(x['personas_con_mismo_rut_y_codigo_v11_y_registro_titulacion']>0 for x in summary),'qa_csv_vs_xlsx':output['qa_csv_vs_xlsx']},ensure_ascii=False))
