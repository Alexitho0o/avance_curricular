import json,hashlib
from pathlib import Path
from collections import Counter,defaultdict
from openpyxl import load_workbook
T=Path('/tmp/ipss_embudo_2317_20261022'); O=Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/auditoria_embudo_2317_20261022')
h=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
r=json.load(open(T/'embudo_privado.json')); q=json.load(open(O/'EMBUDO_AGREGADO.json'))
assert len(r)==4973 and len({x[0] for x in r})==2317
ct=Counter(x[9] for x in r); groups=defaultdict(set)
for x in r:groups[x[9]].add(x[0])
seen=set()
for cause,n,p,e in q['particion']:
 assert ct[cause]==n and len(groups[cause])==p and len(groups[cause]-seen)==e
 seen|=groups[cause]
no=groups['NO_LOCALIZADO_COBERTURA_ACOTADA'];new=json.load(open(T/'evidencias_nuevas.json'))
assert not {x[0] for x in new if x[4]}&no
hashes=json.load(open(T/'hashes.json')); assert all(h(p)==v for p,v in hashes.items())
books=[(O,'AUDITORIA_EMBUDO_2317_IPSS_20261022.xlsx',False),(T,'AUDITORIA_PRIVADA_EMBUDO_2317_IPSS_20261022.xlsx',True)]
checked=0;out={}
for d,f,private in books:
 w=load_workbook(d/f,read_only=False,data_only=False)
 for s in json.load(open(d/'payload.json')):
  expected=([s['headers']]+s['rows']) if private else s['rows']
  actual=w[s['name']]
  for ir,expectedrow in enumerate(expected,1):
   for ic,v in enumerate(expectedrow,1):
    cell=actual.cell(ir,ic)
    assert cell.value==(None if v=='' else v), (s['name'],ir,ic,'value mismatch')
    checked+=1
  if private:
   for ic in s.get('textcols',[]):
    for cells in actual.iter_rows(min_row=2,min_col=ic+1,max_col=ic+1):
     c=cells[0]; assert c.value is None or (isinstance(c.value,str) and c.number_format=='@')
 w.close();out[f]=h(d/f)
prior={'/Users/alexi/Documents/GitHub/avance_curricular/outputs/retencion_puente_provisional_20261022_v6/RETENCION_IPSS_PROVISIONAL_PUENTE_20261022_V6.xlsx':'b1cc8d9e813feb44e548a833a0fd96e1501acc15ce634bd5cf591cb5042fca8b','/Users/alexi/Desktop/REVISION_PRIVADA_CONTRASTE_63_IPSS_20261022_V3.xlsx':'ddfe0244a8f2773942bc858b05d5caae5cc4a8e4b2946e7fccab643524a355df'}
assert all(h(p)==v for p,v in prior.items())
qa={'estado':'PASS','personas':2317,'observaciones':4973,'fuentes_hash_sin_cambios':len(hashes),'celdas_verificadas':checked,'fuera_universo':0,'recuperaciones_nuevas':0,'V6_privadoV3_intactos':True,'archivos_sha256':out}
(O/'QA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa))
