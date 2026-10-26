import json,csv,re,hashlib,unicodedata
from pathlib import Path
from collections import defaultdict,Counter
import openpyxl
R=Path.cwd();T=Path('/tmp/ipss_embudo_2317_20261022');O=R/'outputs/auditoria_embudo_2317_20261022';prior=json.load(open('/tmp/ipss_revision_63_20261022/payload.json'));get=lambda n:next(s['rows'] for s in prior if s['name']==n);states={'CONTRADICCION_O_ATRIBUCION_MULTIPLE','SIN_EVIDENCIA_ANUAL_APLICABLE','CODIGO_ANUAL_SIN_CONTEXTO_DA_EXACTO','DIFIERE_CODIGO_PROPUESTO_OBSERVADO'};units=[r for r in get('CONTRASTE_63') if r[10] in states];persons={r[0] for r in units};assert len(persons)==2317 and len(units)==4973
old=[r for r in get('EVIDENCIAS_SIES') if r[0] in persons];oldhits=defaultdict(set)
for e in old:
 if e[4]=='COMPLETO':oldhits[e[0]].add(e[3])
unknown=persons-set(oldhits);assert len(unknown)==1016
schemas=json.load(open('/tmp/ipss_revision_privada_20261022/schemas_v2.json'));inv=json.load(open(R/'data/cierre_d2/20261016/busqueda_fuentes.json'));inventory=json.load(open('/tmp/ipss_revision_privada_20261022/inventario_sies.json'))
paths=sorted({r['path'] for r in schemas if any(s in Path(r['path']).name for s in ['A_CURRICULAR_2017_162','A_CURRICULAR_2018_162','A_CURRICULAR_2019_162'])})
paths += [r['path'] for r in inv['candidatos'] if Path(r['path']).name in ['MU2026_CIERRE_AUDITORIA_SUBIDAS_REALIZADAS_4165_TRAZABILIDAD.xlsx','MATRICULA_UNIFICADA_PREGRADO_2026_CONSOLIDADA_FINAL_CON_ENCABEZADOS.xlsx','matricula_unificada_2026_COMPLEMENTO_95_CODCLI_PARA_SUBIR.csv','matricula_unificada_2026_pregrado_PARA_SUBIR.csv'] and '/Enviados/Matri' in r['path'] and any(t in r['path'] for t in ['CUARTA SUBIDA','QUINTA SUBIDA','REPORTE JULIO'])]
paths += [p for p in inventory if p.endswith('titulación unificada 2026.xlsx') or Path(p).name=='Titulación Unificada 2026.csv']
paths+=json.load(open(T/'inventario_mu_2022_2025.json'))
new=[];coverage=[];hashes={};body_dvs=defaultdict(set)
def val(v):return '' if v is None else str(int(v)) if isinstance(v,float) and v.is_integer() else str(v).strip()
def norm(v):return unicodedata.normalize('NFKD',val(v)).encode('ascii','ignore').decode().upper().strip()
def key(a,b):
 a=val(a).replace('.','').replace(' ','');b=val(b).upper()
 return str(int(a))+'-'+b if a.isdigit() and re.fullmatch('[0-9K]',b) else None
def valid(k):
 a,b=k.split('-');s=sum(int(n)*(2+i%6) for i,n in enumerate(a[::-1]));v=11-s%11;return b==('0' if v==11 else 'K' if v==10 else str(v))
def rows_xl(p):
 w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 for s in w:yield s.title,s.values
 w.close()
# Encabezado de CSV sin títulos: solo se autoriza si coincide por contenido con filas del XLSX pareado.
titp=next(Path(p) for p in paths if p.endswith('titulación unificada 2026.xlsx'));tw=openpyxl.load_workbook(titp,read_only=True,data_only=True);ts=tw['Res_Tit_Carga'];ti=iter(ts.values);th=list(next(ti));trows=list(ti);normrow=lambda row:tuple(val(v) for v in row);tset={normrow(r) for r in trows};tw.close()
mup=next(Path(p) for p in paths if 'MU2026_CIERRE_AUDITORIA' in p);mw=openpyxl.load_workbook(mup,read_only=True,data_only=True);mi=iter(mw['SUBIDAS_REALIZADAS_4165'].values);mh=list(next(mi));mr=list(mi);mw.close()
def projection(row,h,kind):
 d=dict(zip([norm(v) for v in h],row));k=key(d.get('N_DOC',d.get('NUM_DOCUMENTO')),d.get('DV'));code=('I162S%sC%sJ%sV%s'%tuple(val(d.get(t)) for t in ['COD_SED','COD_CAR','JOR','VERSION'])) if kind=='MU' else val(d.get('COD_SIES_OBT_TIT'));return k,code
ms={projection(r,mh,'MU') for r in mr};ts2={projection(r,th,'TIT') for r in trows}
for path in dict.fromkeys(paths):
 p=Path(path)
 try:
  hashes[path]=hashlib.sha256(p.read_bytes()).hexdigest()
  if p.suffix=='.xlsx':streams=rows_xl(p)
  else:
   raw=p.read_bytes()
   try:text=raw.decode('utf-8-sig')
   except UnicodeDecodeError:text=raw.decode('latin1')
   rr=list(csv.reader(text.splitlines(),delimiter=';'));headerless=False
   if p.name=='Titulación Unificada 2026.csv':
    matched=sum(len(r)==len(th) and projection(r,th,'TIT') in ts2 for r in rr)
    if matched==len(rr):rr=[th]+rr;headerless=True
    else:coverage.append([path,'CSV',0,len(rr),0,'SIN_ENCABEZADO; esquema no validado íntegramente contra XLSX',matched]);continue
   if 'PARA_SUBIR' in p.name and len(rr[0])==len(mh):
    matched=sum(projection(r,mh,'MU') in ms for r in rr)
    if matched==len(rr):rr=[mh]+rr;headerless=True
    else:coverage.append([path,'CSV',0,len(rr),0,'ESQUEMA_MU_NO_VALIDADO',matched]);continue
   streams=[('CSV',iter(rr))]
  for sheet,it in streams:
   header=None;count=match=complete=0;academic='';originalheader=None;sampletypes=Counter()
   for rn,row in enumerate(it,1):
    if header is None:
     hh=[norm(v) for v in row]
     if any(n in hh for n in ['RUN','N_DOC','NUM_DOCUMENTO','RUT']):
      header=hh;originalheader=row
      for cell in row:
       yy=re.search(r'ano academico (20\d{2})',norm(cell).lower())
       if yy:academic=yy.group(1)
     if rn>=15 and header is None:break
     continue
    count+=1;x=dict(zip(header,row));getv=lambda *keys:next((x[k] for k in keys if k in x and x[k] is not None),None)
    k=key(getv('RUN','N_DOC','NUM_DOCUMENTO'),getv('DV'));typ=norm(getv('TIPO_DOC','TIPO_DOCUMENTO'));sampletypes[typ]+=1
    if k and k.split('-')[0] in {v.split('-')[0] for v in unknown}:body_dvs[k.split('-')[0]].add(k.split('-')[1])
    if k not in persons or typ not in ['', 'R','1','RUT']:continue
    match+=1;code=val(getv('CODIGO_UNICO','CODIGO DE LA CARRERA ASOCIADO AL ESTUDIANTE','COD_SIES_OBT_TIT')).upper();field='CODIGO_UNICO_O_TITULO'
    if all(getv(t) is not None for t in ['COD_SED','COD_CAR','JOR','VERSION']):code='I162S%sC%sJ%sV%s'%tuple(val(getv(t)) for t in ['COD_SED','COD_CAR','JOR','VERSION']);field='COMPONENTES_MU'
    full=bool(re.fullmatch(r'I\d+S\d+C\d+J\d+V\d+',code));complete+=full
    proc='AVANCE' if 'A_CURRICULAR' in p.name else 'TITULADOS' if 'titul' in p.name.lower() else 'MATRICULA'
    pathyear=next((v for v in p.parts if v in ['2022','2023','2024','2025','2026']),'2026');yr=academic if academic else pathyear;sem='ANO_ACADEMICO_EXPLICITO_EN_ENCABEZADO' if academic else 'PROCESO_2026; no año de actividad individual' if proc=='TITULADOS' else 'CARGA_MU'+pathyear+'; año por trazabilidad archivo'
    plan=val(getv('PLAN DE ESTUDIOS DE LA CARRERA','PLAN_ESTUDIOS','DA_PLAN_ESTUDIO'));cli=val(getv('CODCLI','DA_CODCLI'));car=val(getv('CODCARPR','DA_CODCARPR'))
    new.append([k,proc,yr,code,'COMPLETO' if full else 'INCOMPLETO',field,plan,cli,car,path,sheet,rn,sem,'DERIVADO_TRAZABILIDAD' if 'TRAZABILIDAD' in p.name else 'ARCHIVO_ENVIADOS_NO_COMPROBANTE_ACEPTACION' if 'Enviados' in path else 'PUBLICADOS_CON_ENCABEZADO_DOCUMENTADO',val(getv('RUN','N_DOC','NUM_DOCUMENTO')),val(getv('DV'))])
   coverage.append([path,sheet,count,match,complete,'LEIDO_CON_ALIAS_RUN' if header and 'RUN' in header else 'LEIDO' if header else 'SIN_TABLA_INDIVIDUAL_RECONOCIDA',json.dumps(dict(sampletypes))])
 except Exception as ex:coverage.append([path,'',0,0,0,'INACCESIBLE_O_ERROR',type(ex).__name__])
# Resumen exclusivo de nuevas localizaciones y variantes de llave; sin promover coincidencias de cuerpo sin DV.
nh=defaultdict(set)
for e in new:
 if e[4]=='COMPLETO':nh[e[0]].add(e[3])
recovered=unknown&set(nh);still=unknown-set(nh)
summary={'personas_foco':2317,'observaciones_foco':4973,'antes_personas_sin_codigo':1016,'nuevas_personas_con_codigo':len(recovered),'restantes_sin_codigo_en_fuentes_incorporadas':len(still),'archivos_dirigidos':len(hashes),'evidencias_nuevas':len(new),'personas_DV_invalido_foco':sum(not valid(k) for k in persons),'personas_DV_invalido_antes_sin_codigo':sum(not valid(k) for k in unknown),'personas_antes_sin_codigo_mismo_cuerpo_otro_DV_en_nuevas_fuentes':sum(any(d!=k.split('-')[1] for d in body_dvs[k.split('-')[0]]) for k in unknown)}
(T/'evidencias_nuevas.json').write_text(json.dumps(new,ensure_ascii=False));(T/'cobertura.json').write_text(json.dumps(coverage,ensure_ascii=False));(T/'hashes.json').write_text(json.dumps(hashes,ensure_ascii=False));(T/'resumen.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2));(O/'RECUPERACION_AGREGADA.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
for p,h in hashes.items():assert hashlib.sha256(Path(p).read_bytes()).hexdigest()==h
print(json.dumps(summary));print('cobertura',[(Path(r[0]).name,r[1],r[2],r[3],r[5]) for r in coverage])
