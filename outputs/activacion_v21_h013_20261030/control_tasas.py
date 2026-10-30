"""Recalcula los numeradores y denominadores con todas las filas de las fuentes aprobadas."""
from collections import defaultdict
from fractions import Fraction
from pathlib import Path
import openpyxl
import json
from construir_v21 import ROOT,OUT,load,sha
from common.analisis_desercion.h001 import reproducir_h001,normalizar
from common.diccionario_sies_seguimiento import rut_normalizado

def rates(api):
    # Índice contextual del puntero verificado. Ningún filtro de cobertura o de código.
    context=defaultdict(set)
    for person,y,period,raw in api.db.execute('select persona,anio,periodo,payload from observaciones'):
        r=json.loads(raw)
        code=r.get('codigo_documental_atribuible') or r.get('codigo_interno_adoptado')
        context[person,str(y),str(period)].add(code or 'E')
    def identity(value):
        try:return rut_normalizado(value)
        except ValueError:return None  # Sin identidad completa: conservar el caso en el total.
    def bucket(person,year,period):
        codes=context.get((person,str(year),str(period)),set())
        return next(iter(codes)) if len(codes)==1 else 'VARIOS' if codes else 'SIN_CONTEXTO'
    def grouped(rows):
        groups=defaultdict(lambda:[0,0])
        for k,n,d in rows:
            groups[k][0]+=int(n);groups[k][1]+=int(d)
        return {'numerador':sum(v[0] for v in groups.values()),'denominador':sum(v[1] for v in groups.values()),'grupos_contextuales':len(groups),'sin_contexto':groups.get('SIN_CONTEXTO',[0,0])}
    frozen=load(ROOT/'outputs/retencion_primer_anio_h001_activo_20261028/RETENCION_PRIMER_ANIO_AGREGADO.json')
    src=Path(frozen['fuente_diagnostico']);assert sha(src)==frozen['fuente_sha256']
    wb=openpyxl.load_workbook(src,read_only=True,data_only=True)
    def read(sheet):
        it=wb[sheet].values;header=next(it)
        return [dict(zip(header,row)) for row in it]
    elig=reproducir_h001(read('BASE_RETENCION_MU'),read('01_CATALOGO_CARRERA'));wb.close()
    result={}
    for year,rr in elig.items():
        result['H001_'+str(year)]=grouped((bucket(identity(normalizar(r['N_DOC'])+'-'+normalizar(r['DV'])),year,1),ret,1) for r,ret in rr)
    rule=load(ROOT/'data/reglas_negocio/indicador_intertrimestral/REGLA.json')
    ref=next(r for r in rule['fuentes'] if r['id']=='V24_COPIA_RECALCULADA')
    p=Path(ref['ruta']);assert sha(p)==ref['sha256']
    wb=openpyxl.load_workbook(p,read_only=True,data_only=True)
    # R19 fija la foto y sus banderas Q/R/T. El recálculo institucional
    # reagrupa los casos por el código consultado, preservando la unidad.
    origin=Path.home()/'.local/share/ipss/aplicacion_diccionario_intertrimestral_20261023_125556/ORIGENES_T1.json'
    assert sha(origin)=='9ba34a21ce89170f6343e068ba38ac1a546a927c1bd5ffcab9f4fd6ef48288a0'
    people={int(o['fila_foto']):o['persona'] for o in json.loads(origin.read_text()) if int(o['anio'])==2026}
    rr=[(i,r) for i,r in enumerate(wb['20_Foto_BBDD'].iter_rows(min_row=6,max_row=4478,max_col=22,values_only=True),6) if r[1]==2026]
    for tag,pred in [('T1_T3',lambda r:r[2]==1),('TOTAL_T3',lambda r:r[19]==1)]:
        result['R19_'+tag]=grouped((bucket(people.get(i) if r[2]==1 else identity(str(r[4])),2026,3),r[16]==1,r[17]==1) for i,r in rr if pred(r))
    # Hoja agregada del área. La unidad de T2 es matrícula, sin identificador individual.
    area=wb['23_Foto_Renuncias_T2_2026'].iter_rows(min_row=6,max_row=24,max_col=6,values_only=True)
    result['R19_INGRESO_T2']=grouped((r[1],(r[2] or 0)-(r[3] or 0),r[2] or 0) for r in area)
    wb.close()
    refs={'H001_'+str(r['cohorte']):(r['retenidos'],r['casos']) for r in frozen['grupos']['institucional']}
    refs.update({'R19_T1_T3':(1360,1467),'R19_TOTAL_T3':(2290,2519),'R19_INGRESO_T2':(462,528)})
    for k,(n,d) in refs.items():
        z=result[k];num,den=z['numerador'],z['denominador'];delta=100*(Fraction(num,den)-Fraction(n,d))
        z.update(tasa_pct=100*num/den,delta_pp=float(delta),referencia=[n,d])
        if (num,den)!=(n,d):raise ValueError(f'{k}: diferencia de numerador/denominador: {(num,den)} frente a {(n,d)}')
    return result
