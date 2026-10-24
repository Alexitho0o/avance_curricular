from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
import copy,hashlib,json,shutil
root=Path('/Users/alexi/Documents/GitHub/avance_curricular')
src=root/'outputs/aclaracion_tasas_2026_20261023_v16/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V16.xlsx'
out=root/'outputs/pendientes_entrega_20261024_v17';out.mkdir(exist_ok=True)
dst=out/'RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261024_V17.xlsx'
assert not dst.exists()
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(src)=='aa4309934a0ed34740913bfdf2314e97a1a4cb51cbe36b830266df13df4cb37c'
N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda t:'{'+N+'}'+t
with ZipFile(src) as z: original={n:z.read(n) for n in z.namelist()}
with ZipFile('/tmp/retencion_v17/hoja.xlsx') as z: donor={n:z.read(n) for n in z.namelist()}
base=E.fromstring(original['xl/styles.xml']); ds=E.fromstring(donor['xl/styles.xml'])
maps={}
for tag in ['numFmts','fonts','fills','borders','cellStyleXfs','cellXfs']:
 b=base.find(q(tag));d=ds.find(q(tag))
 if d is None: continue
 if b is None:
  b=E.Element(q(tag),count='0');base.insert(0,b)
 mp={}
 for i,item in enumerate(d):
  item=copy.deepcopy(item)
  if tag=='numFmts':
   old=int(item.get('numFmtId'));new=max([163]+[int(c.get('numFmtId')) for c in b])+1;mp[old]=new;item.set('numFmtId',str(new))
  else:mp[i]=len(b)
  if tag in ['cellStyleXfs','cellXfs']:
   for attr,typ in [('fontId','fonts'),('fillId','fills'),('borderId','borders'),('numFmtId','numFmts'),('xfId','cellStyleXfs')]:
    if attr in item.attrib and int(item.get(attr)) in maps.get(typ,{}):item.set(attr,str(maps[typ][int(item.get(attr))]))
  b.append(item)
 b.set('count',str(len(b)));maps[tag]=mp
sheet=E.fromstring(donor['xl/worksheets/sheet1.xml'])
strings=E.fromstring(donor['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in donor else []
for c in sheet.iter(q('c')):
 if 's' in c.attrib:c.set('s',str(maps['cellXfs'][int(c.get('s'))]))
 if c.get('t')=='s':
  si=strings[int(c.find(q('v')).text)];c.remove(c.find(q('v')));c.set('t','inlineStr');inline=E.SubElement(c,q('is'))
  for child in si:inline.append(copy.deepcopy(child))
wb=E.fromstring(original['xl/workbook.xml']);sheets=wb.find(q('sheets'));assert len(sheets)==54
rel=E.fromstring(original['xl/_rels/workbook.xml.rels']);rn=rel.tag.split('}')[0]+'}'
rid='rIdEntregaV17';assert rid not in [x.get('Id') for x in rel]
sheetid=max(int(x.get('sheetId')) for x in sheets)+1
E.SubElement(sheets,q('sheet'),name='55_Pendientes_entrega',sheetId=str(sheetid),attrib={'{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id':rid})
part='xl/worksheets/sheet55.xml';assert part not in original
E.SubElement(rel,rn+'Relationship',Id=rid,Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet',Target='worksheets/sheet55.xml')
ct=E.fromstring(original['[Content_Types].xml']);E.SubElement(ct,ct.tag.split('}')[0]+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
enc=lambda x:E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)
changes={'xl/styles.xml':enc(base),'xl/workbook.xml':enc(wb),'xl/_rels/workbook.xml.rels':enc(rel),'[Content_Types].xml':enc(ct),part:enc(sheet)}
with ZipFile(dst,'w',ZIP_DEFLATED) as z:
 for n,v in original.items():z.writestr(n,changes.get(n,v))
 z.writestr(part,changes[part])
with ZipFile(dst) as z:
 assert z.testzip() is None
 for n,v in original.items():
  if n not in changes:assert z.read(n)==v,n
 for n in z.namelist():
  if n.endswith('.xml') or n.endswith('.rels'):E.fromstring(z.read(n))
for tag in ['numFmts','fonts','fills','borders','cellStyleXfs','cellXfs','dxfs']:
 old=E.fromstring(original['xl/styles.xml']).find(q(tag));new=base.find(q(tag))
 if old is not None:
  for i,c in enumerate(old):assert E.tostring(c)==E.tostring(new[i]),tag
import openpyxl
w=openpyxl.load_workbook(dst,read_only=True,data_only=False)
assert len(w.sheetnames)==55 and w.sheetnames[-1]=='55_Pendientes_entrega'
assert w.worksheets[-1]['A5'].value.startswith('66 renuncias')
w.close()
desktop=Path('/Users/alexi/Desktop')/dst.name;assert not desktop.exists();shutil.copy2(dst,desktop);assert sha(dst)==sha(desktop)
qa={'clasificacion':'B/C','estado':'OK_PRESERVACION','hojas_anteriores_intactas':54,'graficos_intactos':sum(n.startswith('xl/charts/chart') and n.endswith('.xml') for n in original),'tablas_intactas':sum(n.startswith('xl/tables/table') and n.endswith('.xml') for n in original),'bases_formulas_denominadores':'Bytes originales conservados','sha256_original':sha(src),'sha256_salida':sha(dst),'salida':str(dst),'escritorio':str(desktop),'partes_modificadas':list(changes),'limite':'No se recalculan ni reauditan los cruces anteriores; hoja nueva informativa sin fórmulas.'}
(out/'QA_Y_MANIFEST.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
shutil.copy2('/tmp/retencion_v17/build.mjs',out/'build.mjs');shutil.copy2(__file__,out/'assemble.py')
print(json.dumps(qa,ensure_ascii=False,indent=2))
