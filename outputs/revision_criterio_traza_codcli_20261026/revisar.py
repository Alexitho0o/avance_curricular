"""Contrasta el caso discrepante en respaldos locales; no exporta identificadores."""
import hashlib
import json
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent
PRIVATE = Path('/Users/alexi/.local/share/ipss/revision_21_ramos_4_expedientes_20261026/DETALLE_CUATRO_CASOS_V2.json')
PRIVATE_OUT = Path('/Users/alexi/.local/share/ipss/revision_criterio_traza_codcli_20261026')
FOLDER = Path('/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/Archivos de chat de Microsoft Teams')
PRIMARY = next(FOLDER.glob('1668382*.xlsx'))
SUPPORT = FOLDER / 'titulados.xlsx'
OTHER = FOLDER / 'TITULADOS_3280.xlsx'

def sha(p):
    h = hashlib.sha256()
    with p.open('rb') as f:
        for block in iter(lambda: f.read(1048576), b''):
            h.update(block)
    return h.hexdigest()

def norm(x):
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x or '').strip().upper()

def body(x):
    return norm(x).replace('.', '').split('-')[0]

def read_rows(ws):
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [(i, dict(zip(head, row))) for i, row in enumerate(it, 2)]

hashes = {str(p): sha(p) for p in [PRIMARY, SUPPORT, OTHER]}
cases = json.loads(PRIVATE.read_text())
w = openpyxl.load_workbook(PRIMARY, read_only=True, data_only=False)
da = read_rows(w['Datos_alumnos'])
titles = read_rows(w['Res_Tit_Unificada'])
upload = read_rows(w['Matricula Unificada 2026 subir'])
found = []
for case in cases:
    ident = case['RUT'].split('-')[0]
    t = [row for _, row in titles if body(row['NUM_DOCUMENTO']) == ident]
    u = [row for _, row in upload if body(row['NUM_DOCUMENTO']) == ident]
    assert len(t) == len(u) == 1
    candidates = [(i, row) for i, row in da if 2 <= i <= 9677 and row['RUT'] == u[0]['NUM_DOCUMENTO']]
    if candidates and norm(candidates[0][1]['CODCLI']) != norm(t[0]['TRACE_CODCLI_DATOS_ALUMNOS']):
        found.append((case, t[0], candidates))
assert len(found) == 1
case, trace, candidates = found[0]
assert len(candidates) == 3
selected = [r for _, r in candidates if norm(r['CODCLI']) == norm(trace['TRACE_CODCLI_DATOS_ALUMNOS'])]
assert len(selected) == 1
headers = ['CODCARPR', 'NOMBRE_L', 'JORNADA', 'ANOINGRESO', 'PERIODOINGRESO',
           'ANOMATRICULA', 'PERIODOMATRICULA', 'ESTADOACADEMICO', 'SITUACION']
public_candidates = []
for i, row in candidates:
    roles = []
    if i == candidates[0][0]: roles.append('primera_coincidencia_rut')
    if norm(row['CODCLI']) == norm(trace['TRACE_CODCLI_DATOS_ALUMNOS']): roles.append('seleccion_traza')
    public_candidates.append({'roles': roles, **{k: norm(row[k]) for k in headers},
                              'ingreso_actual_y_origen_coinciden': all([
                                  norm(row['ANOINGRESO']) == norm(trace['ANIO_INGRESO_CARRERA_ACTUAL']),
                                  norm(row['ANOINGRESO']) == norm(trace['ANIO_INGRESO_CARRERA_ORIGEN']),
                                  norm(row['PERIODOINGRESO']) == norm(trace['SEM_INGRESO_CARRERA_ACTUAL']),
                                  norm(row['PERIODOINGRESO']) == norm(trace['SEM_INGRESO_CARRERA_ORIGEN'])]),
                              'jornada_coincide_traza': norm(row['JORNADA']) == norm(trace['JORNADA'])})
w.close()

ident = case['RUT'].split('-')[0]
check_digit = case['RUT'].split('-')[1]
w = openpyxl.load_workbook(SUPPORT, read_only=True, data_only=False)
backup_da = [(i, r) for i, r in read_rows(w['Hoja1']) if body(r['RUT']) == ident]
backup_titles = [(i, r) for i, r in read_rows(w['Hoja2']) if body(r['NUM_DOCUMENTO']) == ident and norm(r['DV']) == check_digit]
assert len(backup_da) == 2 and len(backup_titles) == 1
assert norm(backup_da[0][1]['CODCLI']) == norm(trace['TRACE_CODCLI_DATOS_ALUMNOS'])
backup = {'hoja1_candidatos': len(backup_da), 'primera_coincidencia_es_candidato_traza': True,
          'carreras': [norm(r['CODCARPR']) for _, r in backup_da],
          'hoja2_titulos_con_documento_y_dv': len(backup_titles),
          'busqueda_auxiliar': 'VLOOKUP por documento, primer resultado exacto en Hoja1',
          'columnas_resultado': ['NOMBRE_L', 'JORNADA', 'ANOINGRESO', 'PERIODOINGRESO']}
w.close()

w = openpyxl.load_workbook(OTHER, read_only=True, data_only=False)
other_rows = read_rows(w['Hoja1'])
exact = [r for _, r in other_rows if body(r['RUT']) == ident and norm(r['DIG']) == check_digit]
w.close()
assert all(sha(Path(p)) == h for p, h in hashes.items())
result = {'clasificacion': 'B observación, C búsqueda, E atribución y fórmula',
          'proceso': 'titulación 2026 / título observado 2025',
          'fuentes': [{'ruta': p, 'sha256': h} for p, h in hashes.items()],
          'trace': {'puntaje': trace['TRACE_SCORE_CANDIDATO'],
                    'razones': trace['TRACE_REASONS_CANDIDATO'],
                    'codigo_sies': trace['COD_SIES_OBT_TIT'],
                    'titulo': trace['NOMBRE_TITULO'],
                    'anio_ingreso_actual_y_origen': norm(trace['ANIO_INGRESO_CARRERA_ACTUAL']),
                    'periodo_ingreso_actual_y_origen': norm(trace['SEM_INGRESO_CARRERA_ACTUAL'])},
          'candidatos': public_candidates, 'respaldo_titulados': backup,
          'otro_respaldo_coincidencias_rut_dv': len(exact),
          'limites': ['La primera coincidencia y los motivos no prueban pesos ni desempate del generador.',
                      'Dos expedientes de la misma carrera comparten jornada y año/período de ingreso.',
                      'Estado actual no acredita fecha efectiva de cambio ni plan formal del título.',
                      'La ausencia en un archivo no prueba ausencia en todo el histórico.']}
payload = json.dumps(result, ensure_ascii=False, indent=2) + '\n'
for c in cases:
    assert c['RUT'] not in payload and c['RUT'].split('-')[0] not in payload
    for x in c['datosalumnos']:
        assert norm(x['CODCLI']) not in payload
PRIVATE_OUT.mkdir(mode=0o700, parents=True, exist_ok=True)
target = PRIVATE_OUT / 'RESULTADO_CASO_SIN_IDENTIFICADORES.json'
target.write_text(payload)
target.chmod(0o600)
print(json.dumps({'candidatos': len(candidates), 'hoja1_titulados': len(backup_da),
                  'coincidencias_titulos_3280': len(exact), 'hashes_conservados': True}, ensure_ascii=False))
