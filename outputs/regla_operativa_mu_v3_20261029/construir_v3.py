#!/usr/bin/env python3
"""H002m: diagnóstico candidato reproducible; ninguna asignación ni carga."""
import argparse,csv,hashlib,io,json,re,sqlite3,subprocess,sys,unicodedata
from collections import Counter,defaultdict
from pathlib import Path
sys.dont_write_bytecode=True
ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
DATA=ROOT/'data/paridad_codigo_sies/regla_operativa_mu_v3';PRIVATE=Path.home()/'.local/share/ipss/h002m_20261029'
sys.path.insert(0,str(ROOT/'outputs/regla_operativa_mu_v2_20261029'))
import construir_v2 as j
old=j.old
sys.path.insert(0,str(ROOT))
from common.codigo_sies_mu2027 import resolver
sha,read,js,txt,family,pct,table=old.sha,old.read,old.js,old.txt,old.family,old.pct,old.table
LABEL={'1':'REGULAR','3':'CONTINUIDAD'}
MISSING=j.MISSING

def norm(v):return ''.join(c for c in unicodedata.normalize('NFD',txt(v).upper()) if not unicodedata.combining(c))
def num(v):
 try:return int(float(v))
 except (TypeError,ValueError):return None

def cohort_key(values):
 return tuple(num(v) for v in values)

def vig(v):
 n=norm(v)
 return {'VIGENTE CON ALUMNOS NUEVOS':'1','VIGENTE SIN ALUMNOS NUEVOS':'2','VIGENTE CON ESTUDIANTES NUEVOS':'1','VIGENTE SIN ESTUDIANTES NUEVOS':'2','NO VIGENTE':'3','1':'1','2':'2','3':'3'}.get(n,'E')

def ctype(v):return {'PLAN REGULAR':'1','PLAN REGULAR DE CONTINUIDAD':'3','PLAN ESPECIAL':'2','1':'1','3':'3','2':'2'}.get(norm(v),'E')
def csvstr(rows):
 if not rows:return ''
 s=io.StringIO(newline='');w=csv.DictWriter(s,fieldnames=list(rows[0]),lineterminator='\n');w.writeheader()
 for r in rows:w.writerow({k:json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(dict,list,set,tuple)) else v for k,v in r.items()})
 return s.getvalue()
def aggregate(rows,fields):
 c=Counter(tuple(r.get(f) for f in fields) for r in rows)
 return [dict(zip(fields,k),filas=n) for k,n in sorted(c.items(),key=lambda kv:str(kv[0]))]
def one(s):return next(iter(s)) if len(s)==1 else None

def classify_type(duration,catalog,formal,support,independent,observed=True):
 """No mayoría: cualquier apoyo contrario impide una propuesta fuerte."""
 matches={t for d,t in catalog if d==duration and t in LABEL} if duration is not None else set()
 ft={t for d,t in formal if t in LABEL};sup=set(support)-{'E'}
 candidates=matches|ft
 conflict=(len(candidates)>1 or (duration is not None and catalog and not matches) or
           (len(candidates)==1 and bool(sup-candidates)))
 if conflict:return 'CONFLICTO','E','E','Fuentes de perfil/duración o apoyo SIES contradictorios'
 t=one(candidates)
 if t and independent>=2:return LABEL[t],t,'C','Dos mediciones independientes concordantes; conversión diagnóstica condicionada'
 if t or sup:return 'DEBIL','E','E','Una fuente de tipo o sólo apoyo del código informado; no desempata'
 return 'E','E','E','Sin evidencia suficiente de tipo; ramos/nombre/prefijo no deciden'

def is_deterministic(rows,field):
 """La etiqueta debe separar varias versiones, con datos completos y >=2 valores."""
 d=defaultdict(set)
 for r in rows:
  value=r.get(field)
  if value in (None,'','E'):return False
  d[value].add(r['code'])
 return len(d)>1 and all(len(v)==1 for v in d.values()) and len(set().union(*d.values()))>1

def same_partition(left,right,field):
 if not is_deterministic(left,field) or not is_deterministic(right,field):return False
 def mapping(rows):
  d=defaultdict(set)
  for r in rows:d[txt(r[field])].add(r['code'])
  return dict(d)
 return mapping(left)==mapping(right)

def predict_registered(group,previous,offers,types,partial=False,strict=False):
 if strict and group=='NUEVO' and partial:return None,'COBERTURA_OFERTA_REGISTRO_PARCIAL_E'
 return old.predict(group,previous,offers,types)

def main(check=False):
 sources=read(OUT/'FUENTES.json');src={e['id']:Path(e['ruta']) for e in sources}
 for s in sources:
  if sha(s['ruta'])!=s['sha256']:raise ValueError('Fuente alterada: '+s['id'])
 built={}
 def put(p,obj):built[p]=obj if isinstance(obj,str) else js(obj)
 def pair(folder,name,rows):put(folder/(name+'.json'),rows);put(folder/(name+'.csv'),csvstr(rows))
 profiles=read(src['outputs/paridad_codigo_sies_20261029/PERFIL_PLANES_H002C.json'])['perfiles']
 pb={(p['codcarpr'],p['CODPESTUD'],p['sede_umas']):p for p in profiles}
 catj=read(src['data/paridad_codigo_sies/regla_operativa_mu_v2/TIPO_PLAN_POR_PLAN_UMAS.json'])['filas']
 links=read(src['enlaces_H002g']);lb={r['fila_fuente_mu']:r for r in links};inputj=read(src['privado_H002j'])
 master=read(src['data/versiones_historicas_sies_umas/fase_5_20261015/maestro.json'])
 hist=defaultdict(list);histby={}
 for r in master['oferta_anual']:
  if r['anio'] in range(2022,2027):
   t=ctype(r['tipo_plan_sies']);v=vig(r['vigencia']);c=r['codigo_unico'];hist[r['anio'],family(c)].append((c,t,v));histby[r['anio'],c]=(t,v)
 # Primarias: sólo columnas necesarias, MRUN temporal en memoria para trayectorias SIES.
 micro=[];forbidden=set()
 for id in ['sies_individual','sies_individual_2026']:
  for raw in csv.DictReader(src[id].open(encoding='utf-8-sig')):
   y=num(raw['cat_periodo'])
   if y not in range(2022,2027):continue
   c=raw['codigo_unico'];pe=raw['mrun'];forbidden.add(pe)
   micro.append({'y':y,'code':c,'fam':family(c),'vig':vig(raw['vigencia_carrera']),'typ':ctype(raw['tipo_plan_carr']),
     'duration':num(raw['dur_total_carr']),'study_duration':num(raw['dur_estudio_carr']),'cohort':(num(raw['anio_ing_carr_act']),num(raw['sem_ing_carr_act'])),
     'origin':(num(raw['anio_ing_carr_ori']),num(raw['sem_ing_carr_ori'])),'form':txt(raw['forma_ingreso']),'semester':num(raw['sem_ing_carr_act']),'pe':pe,'source':id})
 reg=defaultdict(list);sg=defaultdict(list);trajectories=defaultdict(lambda:defaultdict(set))
 for r in micro:
  reg[r['y'],r['code']].append(r);sg[r['y'],r['fam'],r['cohort']].append(r);trajectories[r['pe'],r['fam']][r['y']].add(r['code'])
 transitions=Counter()
 for (_,f),ys in trajectories.items():
  for y,cs in ys.items():
   if len(cs)==len(ys.get(y-1,set()))==1:transitions[y,f,next(iter(ys[y-1])),next(iter(cs))]+=1
 del trajectories
 regoffer=defaultdict(list);regmeta={};vr=[]
 for (y,c),rs in sorted(reg.items()):
  vs={r['vig'] for r in rs};ts={r['typ'] for r in rs};v=one(vs);t=one(ts)
  if v and v!='E' and t and t!='E':regoffer[y,family(c)].append((c,t,v))
  regmeta[y,c]={'vigencias':sorted(vs),'tipos':sorted(ts),'duracion_total':sorted({r['duration'] for r in rs},key=str),'filas':len(rs)}
  vr.append({'anio':y,'codigo_unico':c,'vigencias_registradas':sorted(vs),'vigencia_unica':v if v else 'E','tipo_plan_registrado':sorted(ts),'duracion_total':regmeta[y,c]['duracion_total'],'filas_sies':len(rs),'clase':'B' if len(vs)==1 and 'E' not in vs else 'E','fuente':rs[0]['source'],'extraccion_05062026':histby.get((y,c),('E','E'))[1]})
 pair(DATA,'VIGENCIA_OFICIAL_POR_CODIGO_ANIO',vr)
 # No inferir vigente desde VIG alumno: se usa campo carrera de microdatos.
 public={}
 with src['Matricula_2007_2026_WEB_10_07_2026.csv'].open(encoding='latin1') as f:
  for r in csv.DictReader(f,delimiter=';'):
   y=num(r['AÑO'].removeprefix('MAT_'));c=r['CÓDIGO CARRERA']
   if txt(r['CÓDIGO DE INSTITUCIÓN'])=='162' and y in range(2022,2027):public[y,c]=num(r['TOTAL MATRÍCULA'])
 countmicro=Counter((r['y'],r['code']) for r in micro)
 pubdiff=[{'anio':y,'codigo_unico':c,'sies_registro':countmicro[y,c],'publicacion':public.get((y,c),0)} for y,c in sorted(set(public)|set(countmicro)) if countmicro[y,c]!=public.get((y,c),0)]
 assert not pubdiff,'Registro SIES no cuadra con publicación'
 # Catálogo mallas leído de fuente original: mínimo/máximo, régimen, ramos únicos.
 listing=defaultdict(list)
 import openpyxl
 wb=openpyxl.load_workbook(src['listado'],read_only=True,data_only=True);it=wb['BBDD Bruta'].iter_rows(values_only=True);header=[txt(v) for v in next(it)]
 for line,vals in enumerate(it,2):
  r=dict(zip(header,vals));k=tuple(txt(r.get(f)) for f in ['Codcarr','CODPESTUD','Sede'])
  if k[0] and k[1] and txt(r.get('CODRAMO')):listing[k].append((num(r.get('NIVEL')),txt(r.get('REGIMEN')),txt(r.get('CODRAMO')),line))
 wb.close()
 # Evidencia de duración formal, sin inferir equivalencias nuevas.
 duration=list(csv.DictReader(src['DURACION_ESTUDIOS.tsv'].open(),delimiter='\t'));durby=defaultdict(list)
 for r in duration:
  for c in ({r['CODCARPR_CANONICO']}|set(r['CODCARPR_ALIAS_LIST'].split('|')))-{'','NAN'}:durby[c].append((num(r['DURACION_ESTUDIOS']),txt(r['TIPO_PLAN_CARRERA']),r['CODIGO_UNICO']))
 h112=defaultdict(list);wb=openpyxl.load_workbook(src['h112b_excel'],read_only=True,data_only=True);it=wb['02_H112A_HOJA10'].iter_rows(values_only=True);header=[txt(v) for v in next(it)]
 for line,vals in enumerate(it,2):
  r=dict(zip(header,vals));plan=txt(r.get('CODPESTUD'))
  h112[plan].append({'fila':line,'codigo_unico':txt(r.get('CODIGO_UNICO')),'duracion':num(r.get('DURACION_ESTUDIOS')),'tipo':txt(r.get('TIPO_PLAN_CARRERA')),'ramos':num(r.get('TOTAL_UNIDADES_MEDIDA')),'metodo':txt(r.get('METODO_SELECCIONADO'))})
 wb.close()
 perma=Counter();wb=openpyxl.load_workbook(src['Permanencia.xlsx'],read_only=True,data_only=True);it=wb['Worksheet'].iter_rows(values_only=True);header=[txt(v) for v in next(it)]
 for vals in it:
  r=dict(zip(header,vals));perma[txt(r.get('CARRERA_ALUM')),txt(r.get('COD_SEDE')),num(r.get('TOTAL_RAMOS_MALLA'))]+=1
 wb.close()
 # Apoyo observado: enlace MU→plan gobernado; atributos del código SIES, no vínculo MRUN.
 support=defaultdict(Counter);people26=defaultdict(set);exp26=defaultdict(set);rows26=Counter()
 for l in links:
  sel=l['seleccion']
  if not sel:continue
  k=sel[1],l['plan'],sel[2];y=l['anio'];c=l['codigo_sies'];rm=regmeta.get((y,c),{})
  for t in rm.get('tipos',[]):support[k][y,c,t,tuple(rm['duracion_total'])]+=1
  if y==2026:
   rows26[k]+=1
   if l['identidad']:people26[k].add(tuple(l['identidad']))
   exp26[k].add(txt(sel[0]))
 for l in links:
  if l['identidad']:forbidden.update(txt(v) for v in l['identidad'] if v is not None)
  if l['seleccion']:forbidden.add(txt(l['seleccion'][0]))
 simulated=read(src['outputs/depuracion_codigo_mu2027_20261029/SIMULACION_CATALOGO_NUEVOS_2027_v2.json']);simcount=Counter((r['codcarpr'],r['plan'],r['sede_umas']) for r in simulated)
 cat=[];typeby={};improved=[]
 for prev in catj:
  k=prev['codcarpr'],prev['plan_umas'],prev['sede_umas'];p=pb.get(k,{});raw=listing.get(k,[]);levels=[x[0] for x in raw if x[0] is not None]
  durations={x['semestres_calculados'] for x in p.get('conversion_si_regimen_PROMEDIOS',[]) if x['semestres_calculados'] is not None};d=one(durations)
  formal={(x['duracion'],x['tipo']) for x in h112.get(k[1],[])};dc={(d,t) for d,t,c in durby[k[0]]};su={t for y,c,t,ds in support[k]}
  # H112B no cuenta como fuente independiente de la malla: número de ramos derivado.
  measured=bool(raw) and d is not None
  independent=(1 if measured else 0)+(1 if dc and d is not None and any(d==dd for dd,t in dc) else 0)
  if prev['tipo_plan']=='E':status,t,cl,why=classify_type(d,dc,formal,su,independent,bool(raw))
  else:status,t,cl,why=prev['tipo_plan'],prev['tipo_codigo'],prev['clase'],'HEREDADO_H002j; no reinterpreta los 121 tipos previos'
  ds=sorted(dc,key=str);sp=[{'anio':y,'codigo_unico':c,'tipo_codigo':t,'duracion_total_sies':list(dd),'matriculas_mu_plan':n} for (y,c,t,dd),n in sorted(support[k].items(),key=lambda z:str(z[0]))]
  peers=[{'codcarpr':q['codcarpr'],'plan':q['CODPESTUD'],'tipo_H112B':q['perfil_formal_H112B'],'ramos':q['ramos_distintos']} for q in profiles if q['codcarpr']!=k[0] and k[0] and any(set(g['codigos'])&set(h['codigos']) for g in p.get('perfiles_grupo_gobernado',[]) for h in q['perfiles_grupo_gobernado'])]
  record={'codcarpr':k[0],'plan_umas':k[1],'sede_umas':k[2],'tipo_anterior':prev['tipo_plan'],'clasificacion':status,'tipo_codigo':t,'clase':cl,'estado':'CANDIDATO','motivo':why,'personas_MU2026':len(people26[k]),'expedientes_MU2026':len(exp26[k]),'matriculas_MU2026':rows26[k],'opciones_simuladas_2027':simcount[k],
   'nivel_minimo':min(levels) if levels else None,'nivel_maximo':max(levels) if levels else None,'regimen_listado':sorted({x[1] for x in raw}),'ramos_distintos':len({x[2] for x in raw}),
   'duracion_diagnostica_semestres':d,'conversion_condicionada_PROMEDIOS':p.get('conversion_si_regimen_PROMEDIOS',[]),'nivel_inicial_mayor_1_indicio':min(levels)>1 if levels else None,
   'fuente_1':{'id':'listado','filas':len(raw),'fila_inicial':min(x[3] for x in raw) if raw else None,'limite':'S/A no acredita régimen semestral; conversión de PROMEDIOS es diagnóstico C, no duración certificada'},
   'fuente_2':{'duracion': [{'semestres':dd,'tipo':tt} for dd,tt in ds],'5810':h112.get(k[1],[]),'limite':'H112B reutiliza malla: sus ramos no suman fuente independiente'},
   'fuente_3':[{'total_ramos':n,'expedientes':v,'sin_plan':True} for (c,s,n),v in sorted(perma.items(),key=lambda z:str(z[0])) if c==k[0] and s==k[2]],'fuente_4':sp,
   'contraste_pares_por_familia':peers,'fuentes_independientes_concordantes':independent,'tipos_apoyo':sorted(su),'no_nombre_ni_prefijo':True}
  cat.append(record);typeby[k]=t
  if prev['tipo_plan']=='E':improved.append(record)
 assert len(improved)==102 and len(cat)==223
 pair(DATA,'TIPO_PLAN_POR_PLAN_UMAS_v3',cat)
 # Las cargas se conservan como B/E. Año declarado sólo cuando un archivo/carpeta lo respalda.
 search=read(OUT/'BUSQUEDA_OFERTA.json');loads=defaultdict(list);loadmeta=[]
 for e in search['archivos']:
  p=src[e['id']];rel=e['relativa_enviados_oferta']
  if p.suffix.lower()!='.csv':continue
  raw=p.read_bytes()
  try:text=raw.decode('utf-8-sig');encoding='utf-8-sig'
  except UnicodeDecodeError:text=raw.decode('latin1');encoding='latin1'
  rr=list(csv.reader(text.splitlines(),delimiter=';'));header=rr[0];has_header=header[0].startswith('COD')
  year=2023 if rel.startswith('2022/') and '2023' in p.name else 2024 if rel.startswith('2024/') else 2026 if rel.startswith('2026/') else None
  # No adivinar año 2024 por una carga de noviembre 2023; queda no atribuido.
  width=len(header);parsed=[]
  if has_header and 'VIGENCIA_CARRERA' in header:
   parsed=[dict(zip(header,row)) for row in rr[1:]]
  elif year==2026 and width==51:
   schema=next((src[x['id']] for x in search['archivos'] if x['relativa_enviados_oferta'].endswith('20251010_96470_20241010_13950_OA_2_TP_ARANCELES.csv')),None)
   if schema:
    fields=next(csv.reader(schema.open(encoding='utf-8-sig'),delimiter=';'));parsed=[dict(zip(fields,row)) for row in rr]
  lm={'fuente':e['id'],'archivo':rel,'sha256':e['sha256'],'anio_propuesto':year,'filas':len(rr)-(1 if has_header else 0),'columnas':width,'codificacion':encoding,'vigencia_columna':bool(parsed),'aceptacion':'NO_ACREDITADA_E','clase':'B/E','alcance':'Contraste local, no sustituye registro SIES'};loadmeta.append(lm)
  for row in parsed:
   parts=[txt(row.get(x)) for x in ['COD_SEDE','COD_CARRERA','COD_JORNADA','VERSION']]
   if not all(re.fullmatch(r'\d+',p) for p in parts):continue
   code='I162S'+parts[0]+'C'+parts[1]+'J'+parts[2]+'V'+parts[3]
   if year:loads[year,code].append({'fuente':e['id'],'vigencia':vig(row['VIGENCIA_CARRERA']),'tipo':ctype(row.get('COD_TIPO_PLAN_CARRERA')),'aceptacion':'E'})
 put(OUT/'CARGAS_OFERTA_LOCALIZADAS.json',loadmeta)
 contrasts=[]
 for y,c in sorted(set(histby)|set(reg)|set(loads)):
  hv=histby.get((y,c),('E','E'))[1];rv=one({r['vig'] for r in reg.get((y,c),[])})
  vs=sorted({r['vigencia'] for r in loads.get((y,c),[])})
  state='REGISTRO_SIN_CODIGO_E' if rv is None else 'REGISTRO_MULTIPLE_O_INVALIDO_E' if rv=='E' else 'EXTRACCION_SIN_CODIGO' if hv=='E' else 'COINCIDE_REGISTRO_EXTRACCION' if rv==hv else 'DIFIERE_VIGENCIA_REGISTRADA'
  contrasts.append({'anio':y,'codigo_unico':c,'vigencia_extraccion':hv,'vigencia_registro':rv,'vigencias_cargas_locales':vs,'carga_aceptada':None,'fuentes_cargas':loads.get((y,c),[]),'clasificacion':state,'clase':'B/C/E'})
 pair(DATA,'CONTRASTE_VIGENCIAS',contrasts)
 load_contrast=[]
 for (y,c),rs in sorted(loads.items()):
  for rr in rs:
   rv=one({x['vig'] for x in reg.get((y,c),[])})
   hv=histby.get((y,c),('E','E'))[1]
   load_contrast.append({'anio':y,'codigo_unico':c,'fuente_carga':rr['fuente'],'vigencia_local':rr['vigencia'],'vigencia_registro':rv,'vigencia_extraccion':hv,
    'comparacion_local_registro':'NO_COMPARABLE_E' if rv is None or rv=='E' else 'COINCIDE' if rv==rr['vigencia'] else 'DIFIERE',
    'comparacion_local_extraccion':'NO_COMPARABLE_E' if hv=='E' else 'COINCIDE' if hv==rr['vigencia'] else 'DIFIERE',
    'aceptacion':'E_NO_ACREDITADA','clase':'B/C/E'})
 pair(DATA,'CONTRASTE_CARGAS_LOCALES',load_contrast)
 # Reproduce causa H002j sin ejecutar su main ni escribir antecedentes.
 splitcodes=defaultdict(set);previous=defaultdict(list);mugroups=defaultdict(list)
 for l in links:
  if l['identidad']:previous[tuple(l['identidad']),l['anio'],family(l['codigo_sies'])].append(l)
 def sk(r):return r['anio'],r['familia'],r['codcarpr'],r['plan'],cohort_key(lb[r['fila_mu']]['cohorte_sies']),r['tipo_j']
 for r in inputj:
  if r['tipo_j']!='E' and r['plan'] not in MISSING and lb[r['fila_mu']]['cohorte_sies'][0] is not None:splitcodes[sk(r)].add(r['real'])
  co=cohort_key(lb[r['fila_mu']]['cohorte_sies']);mugroups[r['anio'],r['familia'],co].append(r)
 results=[];diagn253=[];diagn212=[]
 for orig in inputj:
  r=dict(orig);l=lb[r['fila_mu']];sel=l['seleccion'];k=r['codcarpr'],r['plan'],sel[2] if sel else '';t=typeby.get(k,'E');y=r['anio'];f=r['familia'];co=cohort_key(l['cohorte_sies']);offers=regoffer[y,f]
  histcodes={c for c,tt,v in hist[y,f]};regcodes={c for c,tt,v in offers};partial=bool(histcodes-regcodes)
  pred,why=predict_registered(r['grupo'],set(r['previos']),offers,{t} if t in LABEL else set())
  strict,_=predict_registered(r['grupo'],set(r['previos']),offers,{t} if t in LABEL else set(),partial,True)
  typepred,_=old.predict(r['grupo'],set(r['previos']),hist[y,f],{t} if t in LABEL else set())
  r.update(predicho_m=pred,acierto_m=pred==r['real'],predicho_estricto=strict,acierto_estricto=strict==r['real'],tipo_m=t,cobertura_oferta_parcial=partial,motivo_m=why,cohorte=co,predicho_solo_tipo_v3=typepred)
  principal='NO_APLICA';same=False
  if r['grupo']!='OTRO' and not r['acierto']:
   pri=previous.get((tuple(r['identidad']),y-1,f),[]) if r['identidad'] else [];ps={x['plan'] for x in pri};es={x['seleccion'][0] if x['seleccion'] else None for x in pri}
   known=bool(pri) and not (ps&MISSING) and None not in es and r['plan'] not in MISSING and r['expediente'] is not None
   same=known and ps=={r['plan']} and es=={r['expediente']}
   principal,_=j.residual(r,len(splitcodes[sk(r)])>1,same,r['cambio_plan_observado'] or r['cambio_expediente_observado'],hist[y,f],r['tipo_j'])
  r['causa_j']=principal;r['patron_oficial']=False
  if principal=='INCONSISTENCIA_DE_CARGA':
   mu_rs=mugroups[y,f,co];sies_rs=sg[y,f,co];muvec=Counter(x['real'] for x in mu_rs);svec=Counter(x['code'] for x in sies_rs)
   local_same=[x for x in mu_rs if x['plan']==r['plan'] and x['codcarpr']==r['codcarpr']]
   form_mu=[{'code':x['real'],'form':txt(x['for_ing_act']),'semester':lb[x['fila_mu']]['cohorte_sies'][1]} for x in local_same]
   forok=same_partition(form_mu,sies_rs,'form')
   # Semestre sólo puede explicar reparto al agregar la cohorte anual, nunca dentro del mismo semestre.
   year_mu=[x for x in inputj if x['anio']==y and x['familia']==f and x['plan']==r['plan'] and x['codcarpr']==r['codcarpr'] and lb[x['fila_mu']]['cohorte_sies'][0]==co[0]]
   year_sies=[x for x in micro if x['y']==y and x['fam']==f and x['cohort'][0]==co[0]]
   semok=same_partition([{'code':x['real'],'semester':lb[x['fila_mu']]['cohorte_sies'][1]} for x in year_mu],year_sies,'semester')
   exact=muvec==svec and len(muvec)>1
   changed=bool(same and r['grupo']=='ANTIGUO' and r['predicho'] and transitions[y,f,r['predicho'],r['real']]>0)
   # Coincidencia por transición es soporte agregado, no identificación de las personas internas.
   official=exact or changed
   if typepred==r['real'] and t!=r['tipo_j']:state,cause='EXPLICADA','TIPO_PLAN_V3'
   elif forok:state,cause='EXPLICADA','FORMA_INGRESO_PARTICION_AGREGADA_C'
   elif semok:state,cause='EXPLICADA','SEMESTRE_PARTICION_AGREGADA_C'
   elif r['cambio_plan_observado'] or r['cambio_expediente_observado']:state,cause='EXPLICADA','CAMBIO_PLAN_EXPEDIENTE_OBSERVADO'
   elif r['acierto_m']:state,cause='EXPLICADA','VIGENCIA_REGISTRADA'
   elif official:state,cause='INCONSISTENCIA_OFICIAL_REGISTRADA','PATRON_AGREGADO_SIN_EXPLICACION; NO_IDENTIDAD_NI_ERROR_PERSONAL_PROBADO'
   else:state,cause='E_NO_ACREDITADA_EN_REGISTRO','CONTEOS_NO_IDENTICOS_O_SIN_TRANSICION_OFICIAL'
   r['patron_oficial']=state=='INCONSISTENCIA_OFICIAL_REGISTRADA'
   diagnostic={'anio':y,'familia':f,'grupo':r['grupo'],'plan':r['plan'],'codigo_informado':r['real'],'codigo_H002j':r['predicho'],'cohorte_anio':co[0],'cohorte_semestre':co[1],
     'estado':state,'causa':cause,'conteos_MU_familia_cohorte':dict(sorted(muvec.items())),'conteos_SIES_familia_cohorte':dict(sorted(svec.items())),
     'reparto_igual':exact,'transicion_SIES_conteo':transitions[y,f,r['predicho'],r['real']] if r['predicho'] else 0,'tipo_v3':t,'forma_separa_ambos':forok,'semestre_separa_ambos':semok,
     'cambio_plan_observado':r['cambio_plan_observado'],'cambio_expediente_observado':r['cambio_expediente_observado'],'vigencia_codigo_registrada':regmeta.get((y,r['real']),{}).get('vigencias',[]),'clase':'C/E'}
   diagn253.append(diagnostic)
  if principal=='OFERTA_HISTORICA_NO_REFLEJA_CARGA':
   hv=histby.get((y,r['real']),('E','E'))[1];rv=one({x['vig'] for x in reg.get((y,r['real']),[])})
   cause='REGISTRO_VIGENCIA_1_EXTRACCION_DISTINTA' if rv=='1' and hv!='1' else 'REGISTRO_CONFIRMA_SIN_NUEVOS' if rv in {'2','3'} else 'VIGENCIA_REGISTRADA_E'
   diagn212.append({'anio':y,'familia':f,'codigo_informado':r['real'],'causa':cause,'vigencia_extraccion':hv,'vigencia_registro':rv,'cargas_locales':loads.get((y,r['real']),[]),'acierta_H002m':r['acierto_m'],'clase':'B/C/E'})
  results.append(r)
 assert len(diagn253)==253 and len(diagn212)==212
 def collapse_dicts(rs):
  cc=Counter(json.dumps(r,sort_keys=True,ensure_ascii=False) for r in rs)
  return [dict(json.loads(k),matriculas=n) for k,n in sorted(cc.items())]
 d253=collapse_dicts(diagn253);d212=collapse_dicts(diagn212)
 mu_transition=Counter((x['anio'],x['familia'],x['codigo_H002j'],x['codigo_informado']) for x in diagn253 if x['grupo']=='ANTIGUO')
 transition_rows=[{'anio':y,'familia':f,'codigo_previo':a,'codigo_actual':b,'MU_residuales_H002j':n,'SIES_transiciones':transitions[y,f,a,b],
   'control':'IGUAL_AGREGADO' if n==transitions[y,f,a,b] else 'PATRON_COMPATIBLE_CON_DIFERENCIA' if transitions[y,f,a,b]>0 else 'NO_OBSERVADO_SIES_E'} for (y,f,a,b),n in sorted(mu_transition.items())]
 pair(OUT,'TRANSICIONES_AGREGADAS_MU_SIES',transition_rows)
 put(OUT/'DIAGNOSTICO_253.json',{'total':253,'casos_agregados':d253,'resumen':aggregate(diagn253,['anio','familia','estado','causa']),'limite':'Repartos/cambios SIES se comparan sólo agregados. Patrón oficial registrado no prueba error individual. No se enlazan MRUN y documentos.'})
 put(OUT/'DIAGNOSTICO_212.json',{'total':212,'casos_agregados':d212,'resumen':aggregate(diagn212,['anio','familia','causa']),'aceptacion_cargas_locales':'E, no se usa para resolver'})
 # Tablas de forma y semestre sin identificadores, tanto SIES como MU.
 forms=Counter()
 for r in micro:forms['SIES',r['y'],r['fam'],r['code'],r['cohort'][0],r['cohort'][1],r['form']]+=1
 for r in inputj:
  co=lb[r['fila_mu']]['cohorte_sies'];forms['MU',r['anio'],r['familia'],r['real'],co[0],co[1],txt(r['for_ing_act'])]+=1
 pair(DATA,'FORMAS_SEMESTRES_POR_CODIGO',[dict(zip(['origen','anio','familia','codigo','cohorte_anio','semestre','forma_ingreso'],k),matriculas=n) for k,n in sorted(forms.items(),key=lambda z:str(z[0]))])
 # Denominadores intactos; sensibilidad separada no sustituye el indicador principal.
 def metrics(fields,filterfun=lambda r:True):
  gr=defaultdict(list)
  for r in results:
   if filterfun(r):gr[tuple(r[f] for f in fields)].append(r)
  out=[]
  for k,rs in sorted(gr.items()):
   eligible=rs[0]['grupo']!='OTRO' if 'grupo' in fields else True
   hits=[sum(r[x] for r in rs) for x in ['acierto_i','acierto','acierto_m','acierto_estricto']];n=len(rs);flags=sum(r['patron_oficial'] for r in rs);rem=[r for r in rs if not r['patron_oficial']]
   row=dict(zip(fields,k),filas_mu=n,aciertos_i=hits[0],aciertos_j=hits[1],aciertos_m=hits[2],acierto_i_pct=pct(hits[0],n) if eligible else None,acierto_j_pct=pct(hits[1],n) if eligible else None,acierto_m_pct=pct(hits[2],n) if eligible else None,
     sin_prediccion_m=sum(r['predicho_m'] is None for r in rs),familia_cobertura_parcial_filas=sum(r['cobertura_oferta_parcial'] and r['grupo']=='NUEVO' for r in rs),acierto_sensibilidad_cobertura_completa_pct=pct(hits[3],n) if eligible else None,
     patrones_oficiales_sin_explicacion=flags,denominador_sensibilidad_sin_patron=len(rem),acierto_sensibilidad_sin_patron_pct=pct(sum(r['acierto_m'] for r in rem),len(rem)) if eligible else None)
   out.append(row)
  return out
 bygroup=metrics(['anio','grupo']);totals=metrics(['anio'],lambda r:r['grupo']!='OTRO');five=metrics(['anio','familia'],lambda r:r['grupo']!='OTRO' and r['familia'] in j.FIVE);fivegroup=metrics(['anio','familia','grupo'],lambda r:r['grupo']!='OTRO' and r['familia'] in j.FIVE)
 baseline=read(src['data/paridad_codigo_sies/regla_operativa_mu_v2/METRICAS_H002i_H002j.json'])
 assert [(r['anio'],r['grupo'],r['filas_mu'],r['aciertos_i'],r['aciertos_j']) for r in bygroup]==[(r['anio'],r['grupo'],r['filas_mu'],r['aciertos_H002i'],r['aciertos_H002j']) for r in baseline['por_grupo']]
 db=sqlite3.connect(':memory:');db.execute('create table r(y int,g text,hit int)');db.executemany('insert into r values(?,?,?)',[(r['anio'],r['grupo'],int(r['acierto_m'])) for r in results]);assert db.execute('select y,g,count(*),sum(hit) from r group by y,g order by y,g').fetchall()==[(r['anio'],r['grupo'],r['filas_mu'],r['aciertos_m']) for r in bygroup];db.close()
 pair(DATA,'METRICAS_H002i_H002j_H002m',bygroup);pair(DATA,'METRICAS_TOTALES',totals);pair(DATA,'CINCO_FAMILIAS',five);pair(DATA,'CINCO_FAMILIAS_GRUPO',fivegroup)
 # Contraste 2027: motor puro existente; no modifica estado operativo.
 ht=read(src['data/paridad_codigo_sies/20261029/TABLA_PARIDAD_2026_2027.json'])['filas'];offer27=[r for r in ht if r['anio']==2027]
 pending={r['codigo_unico'] for r in offer27 if r['estado']=='NUEVO_2027_SIN_PLAN_E'}
 sim=[]
 for r in simulated:
  f={'anio':2027,'cohorte':2027,'via_solicitada':'NUEVO','familia':r['familia'],'codcarpr':r['codcarpr'],'plan_umas':r['plan'],'sede_umas':r['sede_umas']}
  ans=resolver(f,typeby,offer27,{},pending);declared={x['codigo_unico'] for x in offer27 if family(x['codigo_unico'])==r['familia'] and x['version_para_nuevos_2027']=='sí'}
  comp='NO_COMPARABLE_E' if ans['codigo'] is None else 'COINCIDE' if ans['codigo'] in declared else 'DIFIERE'
  sim.append({k:r[k] for k in ['familia','codcarpr','plan','sede_umas']}|{'via_anterior':r['via'],'via_candidata':ans['via'],'codigo_candidato':ans['codigo'],'motivo':ans['motivo'],'comparacion_H002':comp,'resuelve_revision_real':r['via']=='REVISION_MANUAL' and ans['via']=='NUEVO','clase':'C/E','unidad':'OPCION_CATALOGO_NO_PERSONA'})
 assert sum(r['via_anterior']=='REVISION_MANUAL' for r in sim)==34
 pair(DATA,'SIMULACION_NUEVOS_2027',sim)
 changes=[]
 for r in results:
  if r['grupo']=='OTRO' or r['predicho_m']==r['predicho']:continue
  c='CATALOGO_TIPO_V3' if r['predicho_solo_tipo_v3']!=r['predicho'] else 'VIGENCIA_DISTINTA' if any(histby.get((r['anio'],c),('E','E'))[1]!=v for c,t,v in regoffer[r['anio'],r['familia']]) else 'RECORTE_CODIGOS_SIN_MATRICULA_NO_CORRECCION_VIGENCIA'
  changes.append({'anio':r['anio'],'familia':r['familia'],'grupo':r['grupo'],'causa':c,'acierto_antes':r['acierto'],'acierto_ahora':r['acierto_m']})
 pair(OUT,'CAMBIOS_PREDICCION_H002j_H002m',aggregate(changes,['anio','familia','grupo','causa','acierto_antes','acierto_ahora']))
 missing2026=Counter()
 for l in links:
  if l['anio']!=2026:continue
  if not l['seleccion']:missing2026['SIN_EXPEDIENTE_ENLAZADO']+=1
  else:
   k=l['seleccion'][1],l['plan'],l['seleccion'][2]
   if k not in typeby:missing2026['PLAN_NO_CATALOGADO_O_NO_OBSERVADO']+=1
   else:missing2026['TIPO_CATALOGADO_'+typeby[k]]+=1
 summary={'fecha_ejecucion':'2026-10-30','fecha_decision':'2026-10-29','estado':'CANDIDATO_NO_ACTIVO','clasificacion_102':dict(Counter(r['clasificacion'] for r in improved)),'cobertura_MU2026':dict(missing2026),'cargas_locales_contraste':aggregate(load_contrast,['anio','comparacion_local_registro','comparacion_local_extraccion']),'102_con_MU2026':dict(Counter(r['clasificacion'] for r in improved if r['personas_MU2026'])),'102_con_opciones2027':dict(Counter(r['clasificacion'] for r in improved if r['opciones_simuladas_2027'])),
  '102_matriculas_MU2026':{s:sum(r['matriculas_MU2026'] for r in improved if r['clasificacion']==s) for s in sorted({r['clasificacion'] for r in improved})},
  'vigencia_contraste':aggregate(contrasts,['anio','clasificacion']),'vigencia_no_unica':[r for r in vr if r['clase']=='E'],
  '253_estados':dict(Counter(r['estado'] for r in diagn253)),'253_causas':dict(Counter(r['causa'] for r in diagn253)),'212_causas':dict(Counter(r['causa'] for r in diagn212)),
  'totales':totals,'por_grupo':bygroup,'simulacion':dict(Counter(r['comparacion_H002'] for r in sim)),'revision_34_resueltas':sum(r['resuelve_revision_real'] for r in sim),
  'meta_98_2025_2026':{str(r['anio']):r['acierto_m_pct']>=98 for r in totals if r['anio'] in (2025,2026)},'errores_individuales_oficiales_certificados':0}
 put(OUT/'RESULTADOS.json',summary)
 decisions=sorted([r for r in improved if r['clasificacion'] in {'DEBIL','CONFLICTO','E'} and (r['personas_MU2026'] or r['opciones_simuladas_2027'])],key=lambda r:(-r['personas_MU2026'],-r['opciones_simuladas_2027'],r['plan_umas']))
 md='# H002m · Decisiones pendientes\n\nC/E. No activar catálogo ni cambiar códigos. Personas MU2026 y opciones de catálogo 2027 son unidades distintas; no sumarlas.\n'
 for i,r in enumerate(decisions,1):
  md+=f"\n## {i}. {r['codcarpr']} · {r['plan_umas']} · {r['sede_umas']}\n\n{r['clasificacion']}: {r['personas_MU2026']} personas MU2026; {r['opciones_simuladas_2027']} opciones simuladas. {r['motivo']}. Evidencia: TIPO_PLAN_POR_PLAN_UMAS_v3.json, clave exacta carrera/plan/sede; fuentes 1–4 con hashes en MANIFEST.\n\nDecisión Alexi: ____\n"
 md+='\n## Promoción del candidato\n\nRevisar los tipos C propuestos y la evidencia condicionada. La vigencia de un registro matriculado no acredita oferta completa: no promover el uso prospectivo de microdatos como catálogo de oferta. Mantener la regla activa H002l.\n\nDecisión Alexi: ____\n'
 put(OUT/'LISTA_DECISIONES_ALEXI.md',md)
 proposal={'estado':'CANDIDATO_NO_APLICADO','regla_activa_hash':sha(src['data/reglas_negocio/mu2027/REGLA_CODIGO_SIES_MU2027.json']),'cambio_propuesto':'Sólo sustituir fuente de tipo por catálogo v3 tras aprobación; aceptar únicamente tipo_codigo 1/3 y clase B/C. Conservar antiguos y oferta 2027 del consumidor actual. No trasladar vigencia retrospectiva de matrícula a oferta de nuevos.','resueltas_de_34':summary['revision_34_resueltas'],'requiere':'Aprobación expresa de Alexi, revisión de conversiones condicionadas y QA de consumidor; no se modifica función ni regla activa.'}
 put(OUT/'PROPUESTA_CAMBIO_REGLA.json',proposal)
 notes='''\n## Alcance y límites\n\nD: encargo H002m de Alexi, 29-10-2026. B: atributos y conteos observados; C: cruce y diagnóstico; E: evidencia faltante o contradictoria. No hay nueva regla oficial A en esta entrega. Todo es CANDIDATO.\n\nLos 102 son claves carrera/plan/sede del catálogo H002j, no personas. Los 121 tipos ya determinados se conservan como antecedente C; no se reclasifican. El régimen S/A del listado no está acreditado: la conversión condicionada usa el régimen explícito de PROMEDIOS ya documentado por H002c; no certifica la duración curricular. H112B es derivado de mallas: sus ramos no son una segunda medición independiente. Permanencia no tiene plan y sólo aporta contraste por carrera. No se usa nombre, prefijo ni mayoría SIES para forzar tipo.\n\nLa vigencia oficial registrada es un atributo B de carrera, no VIG de persona. El registro SIES 2026 existente sí incluye vigencia (microdato IPSS catalogado); la publicación agregada no. Se conserva toda la oferta del histórico como control de cobertura, pero el cálculo solicitado usa sólo códigos con vigencia registrada. Un código sin alumnos puede estar abierto y no aparecer en microdatos: la sensibilidad de cobertura completa bloquea esos nuevos. El cálculo utiliza información del mismo año que pretende reproducir (retrospectiva); una mejora no acredita poder predictivo ni carga aceptada. Las cargas locales sin comprobante Finalizado no se llaman aceptadas.\n\nLa etiqueta INCONSISTENCIA_OFICIAL_REGISTRADA significa patrón agregado también observado en SIES, no identidad ni error personal certificado. El registro SIES carece de plan UMAS y no se reconstruye vínculo MRUN–RUT. Los casos sin patrón comparable permanecen E, en lugar de declarar un error oficial. El acierto sin patrón es una sensibilidad con denominador explícito, no reemplaza los denominadores principales. No hay errores individuales oficiales certificados que se puedan retirar de la tasa principal.\n\nLos 253 y 212 son universos de residuales H002j; la misma matrícula aparece sólo una vez en cada diagnóstico disjunto. Las tablas de evidencia por cohorte se repiten como respaldo por caso y no se suman. Los 34 pendientes 2027 son opciones, no alumnos. Los 37 códigos sin plan y antiguos H002l permanecen intactos.\n\nReproducir: `python3 outputs/regla_operativa_mu_v3_20261029/construir_v3.py --check`. El modo --check compara bytes sin reescribir salidas. Fuentes verificadas por SHA-256 en FUENTES.json; sólo salidas agregadas en Git.\n'''
 report='# H002m · Tipos y vigencia registrados; candidata\n\nEjecución: 30-10-2026; decisión/encargo: 29-10-2026.\n\n'+table([{'estado':s,'planes':n,'con_MU2026':summary['102_con_MU2026'].get(s,0),'con_opciones2027':summary['102_con_opciones2027'].get(s,0)} for s,n in sorted(summary['clasificacion_102'].items())],['estado','planes','con_MU2026','con_opciones2027'])
 report+='\n## Cobertura MU2026: planes E frente a planes no observados\n\n'+table([{'categoria':k,'matriculas':v} for k,v in sorted(missing2026.items())],['categoria','matriculas'])+'\nLos 102 perfiles E no son las 1.046 matrículas de tipo desconocido. Ninguna de estas matrículas tiene una clave exacta de esos 102; no rellenar plan por carrera.\n'
 report+='\n## Acierto por año y grupo\n\n'+table(bygroup,['anio','grupo','filas_mu','acierto_i_pct','acierto_j_pct','acierto_m_pct'])
 report+='\n## Total evaluable y sensibilidades\n\n'+table(totals,['anio','filas_mu','acierto_i_pct','acierto_j_pct','acierto_m_pct','patrones_oficiales_sin_explicacion','denominador_sensibilidad_sin_patron','acierto_sensibilidad_sin_patron_pct','acierto_sensibilidad_cobertura_completa_pct'])
 report+='\n## Cinco familias\n\n'+table(five,['anio','familia','filas_mu','acierto_i_pct','acierto_j_pct','acierto_m_pct'])
 report+='\n## Diagnóstico de 253 por familia y año\n\n'+table(aggregate(diagn253,['anio','familia','estado','causa']),['anio','familia','estado','causa','filas'])
 report+='\n## Cambios de código de antiguos: control agregado\n\n'+table(transition_rows,['anio','familia','codigo_previo','codigo_actual','MU_residuales_H002j','SIES_transiciones','control'])
 report+='\n## Diagnóstico de 212\n\n'+table(aggregate(diagn212,['anio','familia','causa']),['anio','familia','causa','filas'])
 report+='\n## Vigencia por año\n\n'+table(summary['vigencia_contraste'],['anio','clasificacion','filas'])
 report+='\n## Cargas locales frente al registro y la extracción\n\n'+table(summary['cargas_locales_contraste'],['anio','comparacion_local_registro','comparacion_local_extraccion','filas'])+'\nEl nombre oficial y la carpeta Enviados no acreditan aceptación. La fuente 2023 contiene fecha de carga diciembre 2022; los archivos 15-09/30-11 de la carpeta 2023 tienen año de proceso no acreditado. No se asignan por su fecha.\n'
 report+='\n## Efecto del recorte de oferta\n\n'+table(aggregate(changes,['anio','causa','acierto_antes','acierto_ahora']),['anio','causa','acierto_antes','acierto_ahora','filas'])+'\nToda mejora por retirar códigos sin matrícula es sensibilidad de cobertura, no resolución de la hipótesis de vigencias históricas incorrectas.\n'
 report+='\n## Simulación 2027\n\n'+js({'comparacion':summary['simulacion'],'resueltas_de_34':summary['revision_34_resueltas'],'meta_98':summary['meta_98_2025_2026']})+'\nLa meta no es un criterio de ajuste de etiquetas ni de exclusión.\n'+notes
 put(OUT/'INFORME.md',report);put(DATA/'README.md','# H002m · Catálogo y evaluación candidatos\n'+notes)
 # Esquemas de salida controlados + búsqueda de tokens exactos sólo en memoria.
 privacy=[];tokens={v for v in forbidden if re.fullmatch(r'\d{6,12}[Kk]?',v)}
 for p,raw in built.items():
  for token in set(re.findall(r'(?<![A-Za-z0-9])\d{6,12}[Kk]?(?![A-Za-z0-9])',raw)) & tokens:
   privacy.append({'archivo':p.name,'tipo':'TOKEN_INDIVIDUAL','valor':'OMITIDO'})
  if re.search(r'(?<!\d)\d{1,2}\.\d{3}\.\d{3}-[0-9Kk]\b',raw):privacy.append({'archivo':p.name,'tipo':'DOCUMENTO_FORMATEADO'})
 assert not privacy,'Salida con posible identificador'
 qa={'estado':'PASS','denominadores_iguales_H002i_H002j':True,'filas_MU':len(results),'particion_por_anio':dict(Counter(r['anio'] for r in results)),'sql_independiente':True,'planes_E_revisados':102,'catalogo_total':223,'diagnostico_inconsistencias':len(diagn253),'diagnostico_oferta':len(diagn212),'registro_SIES_publicacion_diferencias':len(pubdiff),'vigencia_no_unica':len(summary['vigencia_no_unica']),'revision_base_opciones':34,'privacidad_hallazgos':privacy,'fuentes_hash_verificadas':len(sources),'no_modifica_regla_activa':True,'estado_propuestas':'CANDIDATO_NO_ACTIVO'}
 put(OUT/'QA.json',qa)
 for p,raw in built.items():
  if check:
   if not p.exists() or p.read_text()!=raw:raise ValueError('No reproducible: '+str(p.relative_to(ROOT)))
  else:p.parent.mkdir(parents=True,exist_ok=True);p.write_text(raw)
 manifest={'hito':'H002m','clase':'B/C/D/E','estado':'CANDIDATO','commit_base':'4bf6cdd6bbf44582094c16cf3faee7f3b9c0b052','fuentes':sources,'componentes':{str(Path(__file__).relative_to(ROOT)):sha(__file__)},'salidas':{str(p.relative_to(ROOT)):hashlib.sha256(raw.encode()).hexdigest() for p,raw in built.items()}}
 m=DATA/'MANIFEST.json'
 if check:assert m.read_text()==js(manifest),'MANIFEST no reproducible'
 else:
  m.write_text(js(manifest));PRIVATE.mkdir(exist_ok=True,mode=0o700);q=PRIVATE/'PRUEBA_H002m.json';q.write_text(js(results));q.chmod(0o600)
 print(js({'QA':qa,'resumen':{k:v for k,v in summary.items() if k in ['clasificacion_102','102_con_MU2026','253_estados','212_causas','totales','simulacion','revision_34_resueltas']}}))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--check',action='store_true');main(ap.parse_args().check)
