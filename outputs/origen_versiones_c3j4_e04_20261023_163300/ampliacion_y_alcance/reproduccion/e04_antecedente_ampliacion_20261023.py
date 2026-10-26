"""Revisión acotada de antecedentes nuevos de E04. Solo lectura; sin adjudicación."""
from pathlib import Path
import csv,json,re,sys,zipfile,hashlib,xml.etree.ElementTree as ET
from openpyxl import load_workbook
sys.path.insert(0,'/Users/alexi/Documents/GitHub/avance_curricular')
from common.diccionario_sies_seguimiento import rut_normalizado
from e04_leer_msg import CFB
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest();read=lambda p:json.loads(Path(p).read_text())
INV=Path('/tmp/e04_antecedentes_inventario.json');inv=read(INV);paths=inv['paths'];pick=lambda name:next(Path(p) for p in paths if Path(p).name==name)
e=next(x for x in read('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023/ERRORES6_DETALLE.json') if x['id']=='E04');person=e['persona'];codcli=e['CODCLI'];body=person.split('-')[0]
sources={};reg=lambda p:sources.setdefault(str(p),sha(p))
def core(p):
 with zipfile.ZipFile(p) as z:
  root=ET.fromstring(z.read('docProps/core.xml'));return {n.tag.split('}')[-1]:n.text for n in root if n.tag.split('}')[-1] in ['created','modified']}
prior=read(R/'outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_manifest.json')
schema=pick('IPSS2025.xlsx');assert sha(schema)==prior['fuentes_sha256'][str(schema)];reg(schema)
w=load_workbook(schema,read_only=True,data_only=False);heads=[str(x) for x in next(w['Hoja1'].iter_rows(max_row=1,values_only=True))];w.close();assert len(heads)==32
csvout=[]
for name in ['IPSSreset.csv','IPSS2025rect.csv','IPSS2025.csv']:
 p=pick(name);reg(p)
 try:text=p.read_text(encoding='utf-8-sig');enc='utf-8-sig'
 except UnicodeDecodeError:text=p.read_text(encoding='cp1252');enc='cp1252'
 vals=list(csv.reader(text.splitlines(),delimiter=';'));matches=[]
 for rn,row in enumerate(vals,1):
  assert len(row)==32;d=dict(zip(heads,row))
  try:key=rut_normalizado(d['N_DOC']+'-'+d['DV'])
  except (TypeError,ValueError):continue
  if key==person:
   assert d['TIPO_DOC']=='R';matches.append({'fila':rn,'metodo':'TIPO_DOC=R + N_DOC + DV exactos, normalizados','valores':{k:d[k] for k in ['COD_SED','COD_CAR','JOR','VERSION','VIG','ANIO_ING_ACT','ANIO_ING_ORI','FECHA_MATRICULA']}})
 csvout.append({'ruta':str(p),'sha256':sources[str(p)],'filas':len(vals),'columnas':32,'encoding':enc,'encabezado':False,'schema_fuente':str(schema),'coincidencias':matches,'fecha_asignacion_version':'No disponible en estos CSV; FECHA_MATRICULA no es fecha de asignación SIES.'})
g=pick('GESTIONSGE_1703.xlsx');reg(g);w=load_workbook(g,read_only=True,data_only=True);s=w['Hoja1'];h=list(next(s.iter_rows(max_row=1,values_only=True)));notes=[]
resumen={12225:'Consulta por inscripción de asignaturas; la nota registra su inscripción y aviso.',12226:'Consulta de beneficio Junaeb; sin asignación de versión SIES.',12227:'Asignatura desaparecida de EVA; la nota registra su carga y corroboración.',12228:'Bienvenida, inducción y caracterización; sin asignación de versión SIES.'}
for rn,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
 keys=[str(int(row[i]) if isinstance(row[i],(int,float)) else row[i]).strip().replace('.','') for i in [0,2]]
 if body not in keys:continue
 assert keys[0]==keys[1]==body and row[4]=='IINF';t=str(row[1])+' '+str(row[3])
 notes.append({'fila':rn,'CODCARR':row[4],'FECHA_COMUNICO':str(row[7]),'tipo_comunicacion':row[11],'resumen_manual_verificado':resumen[rn],'menciones_version_o_plan':[x for x in ['VERSION','VERSIÓN','PLAN','SIES'] if x in t.upper()]})
w.close();assert len(notes)==4
msg=[]
for path in paths:
 p=Path(path)
 if p.suffix!='.msg' or not ('/Matricula/' in path or '/2025/Oferta academica/' in path):continue
 reg(p);texts=CFB(p).texts();txt=next((v for k,v in texts.items() if k.startswith('__substg1.0_1000')),'');assert txt
 msg.append({'ruta':str(p),'sha256':sources[str(p)],'caracteres_cuerpo':len(txt),'anos_mencionados':sorted(set(re.findall(r'\b20\d{2}\b',txt))),'fechas_explicitas_en_cuerpo':[line.strip() for line in txt.splitlines() if re.match(r'^(Enviado el:|Fecha:|Sent:|Date:)',line.strip())],'menciones_relevantes':[term for term in ['VERSI','IINF','INFORMÁTICA','INFORMATICA','PRECARGA'] if term in txt.upper()],'identificador_foco_literal_en_cuerpo':body in txt or codcli in txt,'alcance':'Cuerpo textual leído; adjuntos no ejecutados ni usados para adjudicar.'})
ppt=pick('Carga proceso de matrícula 2025.pptx');reg(ppt)
with zipfile.ZipFile(ppt) as z:
 image=z.read('ppt/media/image1.png');image_sha=hashlib.sha256(image).hexdigest()
 slides=[n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$',n)]
 pptext=[{'parte':n,'texto':[e.text for e in ET.fromstring(z.read(n)).iter() if e.tag.split('}')[-1]=='t' and e.text]} for n in slides]
prom=pick('PROMEDIOSDEALUMNOS_2071.xlsx');reg(prom);w=load_workbook(prom,read_only=True,data_only=False);promh=[]
for s in w:
 hh=[str(v) for v in next(s.iter_rows(max_row=1,values_only=True)) if v is not None]
 promh.append({'hoja':s.title,'max_row':s.max_row,'max_column':s.max_column,'primera_fila':hh,'campo_VERSION_en_primera_fila':'VERSION' in hh})
w.close()
with zipfile.ZipFile(prom) as z:objects=[n for n in z.namelist() if any(x in n.lower() for x in ['connection','query','vba'])]
code_ext={'.sql','.ipynb','.py','.m','.vba','.ps1','.sh','.bas','.cls','.qry','.rqy','.odc','.dqy','.iqy'}
codepaths=[p for p in paths if Path(p).suffix.lower() in code_ext]
result={'proceso':'MU2025, antecedente de VERSION=1 en E04 IINF20241/J1','clasificacion':'B observado; C inspección y cruce; E regla generadora pendiente','inventario':{'fuente':str(INV),'sha256':sha(INV),'raices_y_conteos':inv['runs'],'total_archivos':len(paths),'archivos_con_extension_codigo_consulta':codepaths,'exclusion':'No se releyeron los 17 libros del paso previo; de IPSS2025.xlsx solo se verificó el encabezado de 32 columnas. Se excluyeron respaldos FCU2026.'},'csv_nuevos':csvout,'gestionsge':{'ruta':str(g),'sha256':sources[str(g)],'core_no_fecha_efectiva':core(g),'hoja':'Hoja1','filas_datos':16395,'encabezados':h,'metodo':'Cuerpo RUT exacto en ambas columnas RUT_ALUMNO y CODCARR=IINF; el archivo no contiene DV ni CODCLI. No se inventa DV.','coincidencias_cuerpo_rut':notes,'limite_identidad':'Cuerpo y carrera coinciden; no equivale a comprobar RUT+DV/CODCLI completos.'},'correos_revisados':msg,'comprobante_carga':{'ruta':str(ppt),'sha256':sources[str(ppt)],'core_no_fecha_efectiva':core(ppt),'texto_slides':pptext,'imagen_sha256':image_sha,'observacion_visual_B':{'parte':'ppt/media/image1.png','fecha_visible':'2025-05-08 22:23:44','estado_visible':'Carga completada','registros_visibles':2371,'ID_evento_visible':467501},'limite':'La captura acredita el estado de una carga agregado; no contiene nombre/hash del CSV ni decisión individual sobre VERSION. Igualar 2371 no prueba por sí solo que sea ese archivo.'},'promedios':{'ruta':str(prom),'sha256':sources[str(prom)],'core_no_fecha_efectiva':core(prom),'estructura_primera_fila':promh,'objetos_consulta_conexion_vba':objects,'limite':'No se vuelve a calcular trayectoria. La estructura de Hoja1 carece de campo VERSION SIES; los metadatos son de mayo, no acreditan antecedente previo al 16-abr.'},'QA':{'fuentes_intactas':all(sha(p)==h for p,h in sources.items()),'identidad_csv_tipo_documento_dv':True,'GESTIONSGE_cuerpo_sin_DV_declarado':True,'bases_denominadores_diccionario_modificados':False,'nuevas_reglas_adoptadas':0,'regla_generadora_V1_localizada':False},'fuentes_sha256':sources,'pendiente_E':'No se localizó exportación, consulta ni decisión anterior al lote del 16-abr que explique V1. Los CSV nuevos conservan V1, pero su posición cronológica respecto a Precarga no queda acreditada.','limite_cobertura':'Inventario de dos raíces concretas; inspección nueva acotada de 3 CSV, GESTIONSGE, 16 cuerpos MSG, estructura PROMEDIOS y comprobante PPTX. No ausencia universal, no consulta de plataforma SIES/UMAS ni historiales remotos.'}
text=json.dumps(result,ensure_ascii=False,indent=2);assert person not in text and codcli not in text and body not in text
out=Path('/tmp/e04_antecedente_ampliacion_20261023.json');out.write_text(text+'\n')
print(json.dumps({'csv_filas':{x['ruta'].split('/')[-1]:x['filas'] for x in csvout},'gestionsge_coincidencias_cuerpo':len(notes),'correos':len(msg),'correos_menciones_asignacion':sum(bool(x['menciones_relevantes']) for x in msg),'QA':result['QA'],'output':str(out),'sha256':sha(out)},ensure_ascii=False,indent=2))
