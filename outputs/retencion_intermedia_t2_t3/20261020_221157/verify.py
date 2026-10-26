from pathlib import Path
import json,hashlib,zipfile,shutil,subprocess,math
import xml.etree.ElementTree as ET
import openpyxl

tmp=Path(__file__).parent
d=json.loads((tmp/'profile.json').read_text())
out=Path(d['out'])
p=out/'CONTINUIDAD_IPSS_INGRESOS_2025_T2_T3_EN_MU2026_20261020_221157.xlsx'
sha=lambda x:hashlib.sha256(Path(x).read_bytes()).hexdigest()
# The authoring API's documented chart settings do not expose fixed bounds.
# Change only native chart scaling XML, preserving all workbook cells and objects.
ns='http://schemas.openxmlformats.org/drawingml/2006/chart'
ET.register_namespace('c',ns)
with zipfile.ZipFile(p) as z:
    entries=[(x,z.read(x.filename)) for x in z.infolist()]
new=[];axis_edits=0
for item,content in entries:
    if '/charts/chart' in item.filename and item.filename.endswith('.xml'):
        root=ET.fromstring(content)
        for axis in root.findall('.//{'+ns+'}valAx'):
            scale=axis.find('{'+ns+'}scaling')
            assert scale is not None
            for tag,val in [('max','1'),('min','0')]:
                e=scale.find('{'+ns+'}'+tag)
                if e is None:e=ET.SubElement(scale,'{'+ns+'}'+tag)
                e.set('val',val)
            axis_edits+=1
        content=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    new.append((item,content))
tmpzip=p.with_suffix('.tmp.xlsx')
with zipfile.ZipFile(tmpzip,'w',zipfile.ZIP_DEFLATED) as z:
    for i,b in new:z.writestr(i,b)
tmpzip.replace(p)
assert axis_edits==1
for src in d['sources']:assert sha(src['path'])==src['sha256'],src['id']
mu_manifest=json.loads(Path(d['sources'][3]['path']).read_text())
assert mu_manifest['salidas']['sha256_excel_repositorio']==d['sources'][0]['sha256']
with zipfile.ZipFile(p) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):ET.fromstring(z.read(name))
w=openpyxl.load_workbook(p,data_only=False)
assert len(w.sheetnames)==6
assert len(w['Resumen']._charts)==1
chart=w['Resumen']._charts[0]
assert chart.y_axis.scaling.min==0 and chart.y_axis.scaling.max==1
assert w['Resumen'].data_validations.dataValidation[0].formula1=='ListaCarreras'
errors=[];formulas=0
for s in w:
    for row in s:
        for c in row:
            formulas+=c.data_type=='f'
            if c.data_type=='e':errors.append((s.title,c.coordinate,c.value))
    for t in s.tables.values():
        a,b,c,e=openpyxl.utils.range_boundaries(t.ref)
        assert [s.cell(b,k).value for k in range(a,c+1)]==[x.name for x in t.tableColumns],t.name
w.close()
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
for person,row in zip(d['people'],w['Cruce_revision'].iter_rows(min_row=7,max_row=6+len(d['people']),values_only=True)):
    exp=[person[x] for x in ['orig_rows','match','vig','zero','continue','onlyzero','absent','origin_discrepancy','multiple_careers']]
    assert list(row[5:14])==exp,person['person']
for row in w['Por_carrera'].iter_rows(min_row=7,values_only=True):
    if row[0] is None:continue
    co,career=row[:2];a=[x for x in d['people'] if x['cohort']==co and x['career']==career]
    exp=[len(a),sum(x['continue'] for x in a),sum(x['onlyzero'] for x in a),sum(x['absent'] for x in a)]
    assert list(row[2:6])==exp
    assert math.isclose(row[6],exp[1]/exp[0],rel_tol=1e-12)
for i,co in enumerate(['2025-T2','2025-T3'],8):
    a=d['totals'][co];s=w['Resumen'];assert s.cell(i,2).value==a['students'];assert s.cell(i,3).value==a['continue'];assert math.isclose(s.cell(i,4).value,a['continue']/a['students'],rel_tol=1e-12)
for s in w:
    for row in s:
        for c in row:
            if c.data_type=='e':errors.append((s.title,c.coordinate,c.value))
assert not errors,errors
w.close()
mf=out/'MANIFEST_CONTINUIDAD.json';manifest=json.loads(mf.read_text());manifest['salida']['sha256']=sha(p)
repo='/Users/alexi/Documents/GitHub/avance_curricular'
manifest['git'].update(rama_despues=subprocess.check_output(['git','-C',repo,'branch','--show-current'],text=True).strip(),head_despues=subprocess.check_output(['git','-C',repo,'rev-parse','HEAD'],text=True).strip())
assert manifest['git']['head_antes']==manifest['git']['head_despues']
manifest['qa']={'resultado':'OK','fuentes_hash_intactos':True,'manifest_mu_hash_verificado':True,'formulas':formulas,'errores_formula':0,'filas_persona_verificadas':len(d['people']),'controles_carrera':'OK','recalculo_y_selector':'artifact-tool; pruebas de todas las carreras y perturbación VIG restaurada','zip_xml_tablas':'OK','grafico':'Eje 0%-100% corregido únicamente en XML de gráfico nativo; verificación de lectura independiente','apertura_excel_nativo':'No realizada'}
for name in ['profile.py','build.mjs','verify.py']:
    if (tmp/name).resolve()!=(out/name).resolve():shutil.copy2(tmp/name,out/name)
manifest['scripts']=[{'archivo':name,'sha256':sha(out/name)} for name in ['profile.py','build.mjs','verify.py']]
mf.write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
inspection=out/(p.name+'.inspect.ndjson')
if inspection.exists():inspection.unlink()
print(json.dumps({'salida':manifest['salida'],'qa':manifest['qa'],'git':manifest['git']},ensure_ascii=False,indent=2))
