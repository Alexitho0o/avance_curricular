"""Prepare cell edits and privacy-safe OOXML transport. Artifact Tool owns new cell authoring/recalculation."""
from pathlib import Path
from collections import defaultdict
import json,zipfile,re,copy,hashlib
import openpyxl
from lxml import etree as E
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');O=R/'outputs/retencion_intermedia_t2_t3/20261021_153436';P=R/'.local_restricted/retencion_intermedia/20261021_153436'
d=json.loads((O/'MODELO_AGREGADO.json').read_text());T=Path(d['template']);w=openpyxl.load_workbook(T,data_only=False);wc=openpyxl.load_workbook(T,data_only=True)
old9='9_Base_Analitica_Detalle';old11='11_BASE_RETENCION_MU_RAW';new11='11_Entradas_Agregadas';old10='10_SIES_2025_Granular';new10='10_Seguimiento_2025'
renames={old11:new11,old10:new10};edits=defaultdict(dict);N=len(d['groups'])+6;N11=len(d['groups'])+4
def put(s,c,v):edits[renames.get(s,s)][c]=v
def args(s):
 out=[];start=0;depth=0;quoted=False
 for i,ch in enumerate(s):
  if ch=='"':quoted=not quoted
  if quoted:continue
  if ch=='(':depth+=1
  elif ch==')':depth-=1
  elif ch==',' and depth==0:out.append(s[start:i]);start=i+1
 return out+[s[start:]]
def weighted(f):
 # Rewrite only COUNTIF(S) over the new aggregate source. Other rank counts retain their meaning.
 pat=re.compile(r'\b(COUNTIFS|COUNTIF)\(')
 while True:
  changed=False
  for m in list(pat.finditer(f))[::-1]:
   level=1;q=False;i=m.end()
   for j in range(i,len(f)):
    if f[j]=='"':q=not q
    if q:continue
    if f[j]=='(':level+=1
    elif f[j]==')':
     level-=1
     if level==0:break
   aa=args(f[i:j])
   if old9 not in aa[0]:continue
   weight=f"'{old9}'!$A$7:$A${N}"
   replacement='SUMIFS('+weight+','+','.join(aa)+')' if m.group(1)=='COUNTIFS' else 'SUMIF('+','.join(aa)+','+weight+')'
   f=f[:m.start()]+replacement+f[j+1:];changed=True
  if not changed:break
 return f
def formula(f):
 for a,b in renames.items():f=f.replace(a,b)
 f=f.replace('$2385',f'${N}')
 f=weighted(f)
 pat=re.compile(r'\b(SUMIFS|SUMIF)\(')
 for m in list(pat.finditer(f))[::-1]:
  level=1;q=False;i=m.end()
  for j in range(i,len(f)):
   if f[j]=='"':q=not q
   if q:continue
   if f[j]=='(':level+=1
   elif f[j]==')':
    level-=1
    if level==0:break
  aa=args(f[i:j])
  if old9 not in aa[0]:continue
  if any('$Y$7:' in a for a in aa) and any('$BO$12' in a for a in aa) and any('$BO$13' in a for a in aa):continue
  if m.group(1)=='SUMIF':aa=[aa[2],aa[0],aa[1]]
  extra=[f"'{old9}'!$Y$7:$Y${N}",'1',f"'{old9}'!$E$7:$E${N}","'14_Vistas_Ejecutivas'!$BO$12",f"'{old9}'!$C$7:$C${N}","'14_Vistas_Ejecutivas'!$BO$13"]
  f=f[:m.start()]+'SUMIFS('+','.join(aa+extra)+')'+f[j+1:]
 return f
for s in w:
 for row in s:
  for c in row:
   if c.data_type=='f' and s.title not in [old9,old11]:
    ff=formula(c.value)
    if s.title=='16_Comparacion_Nacional_SIES' and c.column>=3:
     ff=f'=IF(AND(\'14_Vistas_Ejecutivas\'!$L$4="T1 anual MU",\'14_Vistas_Ejecutivas\'!$H$4="Todos",\'14_Vistas_Ejecutivas\'!$S$4="Todas"),{ff[1:]},"")'
    if s.title=='17_Fuentes_Comparacion_SIES' and old9 in c.value:
     ff=f'=IF(\'14_Vistas_Ejecutivas\'!$L$4="T1 anual MU",{ff[1:]},"")'
    put(s.title,c.coordinate,ff)
   elif isinstance(c.value,str) and s.title in ['4_Resultados_Sector','5_Evolucion','12_Desglose_Jornada_Edad_Sexo','13_Granular_Codigo_Unico','14_Vistas_Ejecutivas','15_Tablas_Ejecutivas']:
    text=c.value.replace('Retenidos','Con registro').replace('retenidos','con registro').replace('Tasa de retención','Continuidad (%)').replace('Tasa de   retención','Continuidad (%)')
    if text!=c.value:put(s.title,c.coordinate,text)
def clear_sheet_data(name,minrow):
 for row in w[name].iter_rows(min_row=minrow):
  for c in row:
   if c.value is not None:put(name,c.coordinate,None)
# These two input sheets are physically emptied in the transport scaffold below.
headers9=['Estudiantes (peso C)','Cohorte de ingreso D','Carrera B','Área SIES B','Sector IPSS C','Continuidad seleccionada C','','','','','','','Sexo B','','Modalidad (código B)','Jornada (código B)','Jornada real C','','','Tramo edad C','COD_CAR B','COD_SED B','VERSION B','Código único B/C','Cohorte seleccionada C']
for i,h in enumerate(headers9,1):put(old9,f'{openpyxl.utils.get_column_letter(i)}6',h or None)
mapcol={'A':'N','B':'A','C':'C','D':'D','E':'E','M':'F','O':'G','P':'H','Q':'I','T':'J','U':'K','V':'L','W':'M','X':'B'}
for i,group in enumerate(d['groups'],7):
 rr=i-2
 for dest,source in mapcol.items():put(old9,f'{dest}{i}',f"='{new11}'!{source}{rr}")
 put(old9,f'F{i}',f'=IF(\'14_Vistas_Ejecutivas\'!$L$4="T1 anual MU",\'{new11}\'!O{rr},IF(\'14_Vistas_Ejecutivas\'!$L$4="T2 observado",\'{new11}\'!P{rr},IF(\'14_Vistas_Ejecutivas\'!$L$4="T3 observado",\'{new11}\'!Q{rr},\'{new11}\'!Y{rr})))')
 put(old9,f'Y{i}',f'=IFERROR(INDEX(\'14_Vistas_Ejecutivas\'!$BQ$14:$BQ$17,MATCH(B{i},\'14_Vistas_Ejecutivas\'!$BP$14:$BP$17,0)),0)')
headers11=['COHORTE D','CODIGO_UNICO B/C','CARRERA B','AREA_SIES B','SECTOR C','SEXO B','MODALIDAD B','JOR B','JORNADA_REAL C','TRAMO_EDAD C','COD_CAR B','COD_SED B','VERSION B','ESTUDIANTES D/C','T1_CONTINUIDAD_ANUAL C','T2_MATRICULA_OBSERVADA C','T3_MATRICULA_OBSERVADA C','T2_VIGENTE_EN_EXTRACCION B/C','T3_VIGENTE_EN_EXTRACCION B/C','LLAVE_PENDIENTE E','CON_PERIODO_UMAS C','TITULADO_OBSERVADO B/C','EGRESADO_OBSERVADO B/C']
headers11+=['LOCALIZADO_ULTIMO_RESPALDO B/C','MATRICULA_2026_VIGENTE_ULTIMO_RESPALDO B/C','TITULADO_ULTIMO_RESPALDO B/C','EGRESADO_ULTIMO_RESPALDO B/C']
for j,h in enumerate(headers11,1):put(old11,f'{openpyxl.utils.get_column_letter(j)}4',h)
for i,row in enumerate(d['groups'],5):
 for j,v in enumerate(row,1):put(old11,f'{openpyxl.utils.get_column_letter(j)}{i}','Sin campo' if j==18 and row[0]==2025 else v)
put(old11,'A1','11. Entradas agregadas de las cohortes y su seguimiento')
put(old11,'A3','C: conteos agregados por dimensiones. Detalle de los cruces y fórmulas por persona en el libro privado. Sin identificadores ni fechas de nacimiento en esta entrega.')
put(old9,'A1','9. Base analítica agregada — cohorte y seguimiento seleccionados')
put(old9,'B3','C: una fila representa un grupo por cohorte, código único, sexo y tramo de edad. Estudiantes es el peso del grupo; las vistas suman ese peso. F cambia con el selector T1/T2/T3 de la hoja 14. Sin datos individuales.')
s14='14_Vistas_Ejecutivas'
put(s14,'C4','Cohortes 2024 y 2025');put(s14,'H4','Todos');put(s14,'S4','Todas');put(s14,'L4','Último respaldo')
for i in range(2,17):put(s14,f'BL{i}',None)
for i,value in enumerate(['Cohortes 2024 y 2025','Cohorte 2024','Cohorte 2025'],2):put(s14,f'BL{i}',value)
for i,value in enumerate(['T1 anual MU','T2 observado','T3 observado','Último respaldo'],2):put(s14,f'BK{i}',value)
for r in [14,15]:put(s14,f'BQ{r}',0)
for r,year in [(16,2024),(17,2025)]:put(s14,f'BQ{r}',f'=IF(OR($C$4="Cohortes 2024 y 2025",$C$4="Cohorte {year}"),1,0)')
put(s14,'B3','Elija cohorte, seguimiento (L4), sector y carrera. T1 conserva el cálculo anual. T2/T3 muestra presencia en los listados parciales del año siguiente; su tasa formal sigue pendiente. Último respaldo muestra matrícula declarada 2026 VIGENTE en la extracción nominal 21-oct-2026, sin certificar cierres anteriores. Las ausencias no son desertores. Para reunir grupos, sume numeradores y estudiantes y divida los totales; no promedie porcentajes.')
put(s14,'B34','C: con registro / cohorte fija. Cada modalidad o carrera puede tener un porcentaje distinto. El total pondera por estudiantes. E: T2/T3 requieren cobertura y estado efectivo de cierre para considerarse tasas formales.')
for c in ['G31','L40','R48','X58','AD31','AI31','AN34']:
 put(s14,c,'C: distribución de la continuidad documentada del seguimiento elegido. E: no interpretar falta de registro como deserción ni estado de extracción como estado efectivo del cierre.')
# Replace inherited conclusions that assert T1-specific values.
for row in w[s14].iter_rows(min_row=20,min_col=45,max_col=49):
 for c in row:
  if isinstance(c.value,str) and c.data_type!='f':put(s14,c.coordinate,'C/E: lectura del grupo seleccionado. T2/T3 describe matrícula documentada; las tasas formales y causas de ausencia siguen pendientes.')
put('Portada','B3','Seguimiento de cohortes IPSS — T1, T2 y T3')
put('Portada','B4','Trayectoria de matrícula de las cohortes 2024 y 2025, con la presentación del libro T1 aprobado')
put('Portada','C7','21-10-2026');put('Portada','C8','Cohorte 2024: 558 estudiantes. Cohorte 2025: 1.150. Seguimientos al año siguiente por matrícula anual y períodos T2/T3.')
put('Portada','C10','NO. Este libro contiene exclusivamente agregados. El detalle identificable permanece separado, en carpeta privada excluida de Git.')
put('Portada','C11','Análisis interno IPSS. T2/T3 documentan matrícula observada; no son publicación SIES ni acreditación de deserción.')
put('Portada','B22','Use la hoja 14 para filtrar cohorte, seguimiento T1/T2/T3, sector y carrera. Las hojas 12, 13 y 15 responden a esa misma selección. La hoja 11 contiene entradas agregadas y la 9 las conecta con las vistas. Las referencias nacionales de 16/17 son anuales y solo se comparan en T1. El detalle individual se entrega en otro libro privado.')
put('1_Que_es_Retencion','B19','C/E: en esta adaptación T1 conserva el indicador anual. T2/T3 muestra matrícula documentada en el período, con la misma cohorte de base. Sin acreditación de cobertura/estado al cierre, ese porcentaje no se presenta como tasa oficial de retención ni su complemento como deserción.')
paragraphs={
 '2_Origen_Datos':['2. Origen de los datos','B: matrícula unificada informada a SIES 2022–2026, conservada por hash. La selección anual aprobada fija 558 y 1.150 estudiantes.','B/C: listados T2/T3, respaldos DatosAlumnos 10246/11998/12112 y 58 extracciones nominales Base Maestra. Se cruzan documentos y CODCLI exactos, conservando matrículas y procedencia.','C: fuentes, hashes, criterio, controles y límites en el manifiesto de esta entrega. No se publican identificadores en este libro.','D: no se modifican originales, reglas del motor anual ni denominadores. Los ejercicios de ingresos T2=222 y T3=124 permanecen separados.','E: dos correspondencias CODCLI pendientes; ninguna se elimina del denominador ni se clasifica como deserción.'],
 '3_Como_se_Cruza':['3. Cómo se construyó el seguimiento','C: se toma cada persona de la cohorte anual aprobada. Se buscan todas sus matrículas por documento exacto y correspondencia CODCLI documentada, sin seleccionar una sola carrera ni inferir por nombre.','C: T2 y T3 se asignan por ANOMATRICULA y PERIODOMATRICULA de la fila fuente. La persona cuenta una vez si tiene alguna matrícula de ese período. Estado y fecha de extracción se conservan separados.','C: por persona se reconstruye la secuencia 2024–2026 y la última matrícula observada. Esto mide trayectoria de matrícula; no acredita asignaturas aprobadas ni fecha de abandono.','C: el detalle privado permite filtrar cada trayectoria y ver sus matrículas, fuentes y filas. En este libro se agregan por carrera/sector/modalidad/jornada/sexo/edad. Todas las vistas suman numeradores y denominadores compatibles.','E: TITULADO/EGRESADO observados no cambian automáticamente el denominador. Falta acreditar programa y fecha y aprobar la regla interna correspondiente.'],
 '6_Limitaciones':['6. Límites que permanecen abiertos','E: las fuentes de cierre originales son parciales. Las matrículas de otros respaldos amplían el seguimiento documentado, pero no certifican población ni estado efectivo al cierre.','E: la ausencia en las fuentes revisadas no acredita deserción. La presencia de una matrícula tampoco acredita actividad o avance curricular.','B/C: una persona puede tener varias carreras/CODCLI y filas por fuente. Se conserva todo el detalle, pero el denominador y cada presencia por período cuentan personas una vez.','E: dos CODCLI continúan pendientes. El documento de origen se conserva en el privado y las personas siguen incluidas en su cohorte.','A/C/E: las tasas nacionales SIES son anuales. No se calcula brecha T2/T3 contra ellas. No se cambia el motor anual ni se aplica una regla nueva para graduación.']}
for name,texts in paragraphs.items():
 for row in w[name].iter_rows(min_row=4,min_col=2):
  for c in row:
   if c.value is not None:put(name,c.coordinate,None)
 for r,text in zip([4,5,10,15,20,25],texts):put(name,f'B{r}',text)
put('3_Como_se_Cruza','B10','C: presencia en T2/T3 se acredita por pertenencia al listado de cierre correspondiente. ANOMATRICULA/PERIODOMATRICULA se conserva como matrícula declarada, separado de la fecha del respaldo y del cierre; no se interpreta como último período cursado ni renovación.')
# Cohort-series slots remain in their existing positions; unused historical cohorts are blank.
for name in ['4_Resultados_Sector','5_Evolucion','12_Desglose_Jornada_Edad_Sexo','13_Granular_Codigo_Unico']:
 for row in w[name]:
  for c in row:
   if c.value in [2022,2023,'2022','2023']:put(name,c.coordinate,None)
# Sector and evolution tables: compact 2024/2025 data, preserving original object positions.
for r in range(7,19):
 if r<=12:
  y=2024+(r-7)//3;sec=['Tecnologías','Administración y Comercio','Salud'][(r-7)%3]
  put('4_Resultados_Sector',f'B{r}',y);put('4_Resultados_Sector',f'C{r}',sec)
  for dest,expr in [('D',f'SUMIFS(\'{old9}\'!$A$7:$A${N},\'{old9}\'!$B$7:$B${N},B{r},\'{old9}\'!$E$7:$E${N},C{r})'),('E',f'SUMIFS(\'{old9}\'!$F$7:$F${N},\'{old9}\'!$B$7:$B${N},B{r},\'{old9}\'!$E$7:$E${N},C{r})'),('F',f'IF(D{r}=0,"",E{r}/D{r})')]:put('4_Resultados_Sector',f'{dest}{r}','='+expr)
 else:
  for c in 'BCDEF':put('4_Resultados_Sector',f'{c}{r}',None)
for r,y in [(7,2024),(8,2025),(9,None),(10,None)]:
 put('5_Evolucion',f'B{r}',y)
 for c in 'CDEFG':
  orig=w['5_Evolucion'][f'{c}{r}'].value
  if y is None:put('5_Evolucion',f'{c}{r}',None)
  elif isinstance(orig,str) and orig.startswith('='):put('5_Evolucion',f'{c}{r}',formula(orig))
 if y:
  original_rate=wc['5_Evolucion'][f'F{9 if y==2024 else 10}'].value
  put('5_Evolucion',f'F{r}',f'=IF(\'14_Vistas_Ejecutivas\'!$L$4="T1 anual MU",{original_rate},"")')
  put('5_Evolucion',f'G{r}',f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),(E{r}-F{r})*100,"")')
put('5_Evolucion','A1','5. Continuidad documentada de las cohortes 2024 y 2025')
put('13_Granular_Codigo_Unico','A1','13. Seguimiento por código único — T1, T2 y T3')
put('13_Granular_Codigo_Unico','B3','C: código único y dimensiones de la cohorte original. A compara T2 y T3 de la cohorte 2025; B muestra la continuidad del seguimiento elegido en la hoja 14. Las cifras de SIES anuales permanecen como referencia en 16/17 y no son tasas T2/T3. Las fórmulas suman estudiantes y con registro; no se promedian porcentajes.')
put('13_Granular_Codigo_Unico','B6','A. Cohorte 2025 — comparación de matrículas documentadas en 2026-T2 y 2026-T3')
for c,v in {'I7':'Con registro T2','J7':'Con registro T3','K7':'Diferencia T3 − T2','L7':'Variación documentada (pp)'}.items():put('13_Granular_Codigo_Unico',c,v)
put('13_Granular_Codigo_Unico','B49','B. Seguimiento elegido — cohortes 2024 y 2025 por código único')
codes2025=sorted({r[1] for r in d['groups'] if r[0]==2025});allcodes=sorted({(r[0],r[1]) for r in d['groups']})
clear_sheet_data(old10,6)
for c,val in {'A1':'10. Cohorte 2025 — matrículas documentadas en T2 y T3 de 2026','B3':'B/C: conteos derivados del cruce exacto de las fuentes verificadas. No es una tabla granular SIES ni una tasa de cierre acreditada.','B6':'Código único','C6':'Carrera','D6':'Sector','E6':'Con registro T2','F6':'Con registro T3'}.items():put(old10,c,val)
codeinfo={r[1]:(r[2],r[4]) for r in d['groups']}
for i,code in enumerate(codes2025,7):
 for c,val in [('B',code),('C',codeinfo[code][0]),('D',codeinfo[code][1])]:put(old10,f'{c}{i}',val)
 for dest,col in [('E','P'),('F','Q')]:put(old10,f'{dest}{i}',f'=SUMIFS(\'{new11}\'!${col}$5:${col}${N11},\'{new11}\'!$A$5:$A${N11},2025,\'{new11}\'!$B$5:$B${N11},B{i})')
for r in range(8,44):
 ix=r-8
 if ix<len(codes2025):
  for cc,src in [('B','B'),('C','C'),('D','D')]:put('13_Granular_Codigo_Unico',f'{cc}{r}',f"='{new10}'!{src}{ix+7}")
  put('13_Granular_Codigo_Unico',f'I{r}',f'=IF(G{r}=0,"",\'{new10}\'!E{ix+7})');put('13_Granular_Codigo_Unico',f'J{r}',f'=IF(G{r}=0,"",\'{new10}\'!F{ix+7})')
  put('13_Granular_Codigo_Unico',f'K{r}',f'=IF(G{r}=0,"",J{r}-I{r})');put('13_Granular_Codigo_Unico',f'L{r}',f'=IF(G{r}=0,"",(J{r}-I{r})/G{r}*100)')
 else:
  for cc in 'BCDEFGHIJKL':put('13_Granular_Codigo_Unico',f'{cc}{r}',None)
for r in range(51,136):
 ix=r-51
 if ix<len(allcodes):
  y,code=allcodes[ix];career,sec=codeinfo[code]
  for cc,val in [('B',y),('C',code),('D',career),('E',sec)]:put('13_Granular_Codigo_Unico',f'{cc}{r}',val)
 else:
  for cc in 'BCDEFGHIJ':put('13_Granular_Codigo_Unico',f'{cc}{r}',None)
for c in ['B45','B137','B139','B141','B143','B145','B147','B149','B151','B153','B155','B157','B159','B161','B163','B165','B167','B169','B171','B173','B175','B177','B179','B181','B183']:
 if w['13_Granular_Codigo_Unico'][c].value is not None:put('13_Granular_Codigo_Unico',c,'C/E: matrícula documentada en el período. La falta de registro y el estado observado no acreditan por sí solos deserción ni vigencia efectiva al cierre.')
put('16_Comparacion_Nacional_SIES','B2','A/C/E: estas referencias SIES son anuales. En T1 se presenta el control anual de las cohortes disponibles; en T2/T3 se desactivan las brechas porque sus períodos no son comparables con la publicación anual. Seleccione T1 anual MU en hoja 14 para consultar el control.')
for s in ['12_Desglose_Jornada_Edad_Sexo','15_Tablas_Ejecutivas']:
 put(s,'B3','C: las vistas responden a cohorte, seguimiento, sector y carrera de la hoja 14. Estudiantes es la cohorte fija; Con registro es el numerador documentado. Sume numeradores y denominadores para comparar totales. E: T2/T3 no acredita cierre ni deserción.')
# Apply shared selections to newly introduced aggregate formulas too.
for name,ee in edits.items():
 for addr,v in list(ee.items()):
  if isinstance(v,str) and v.startswith('=') and old9 in v:ee[addr]=formula(v)
# Remove annual comparisons for years outside this analysis; prevents empty-string arithmetic.
for r in [8,9]:
 for cc in 'CDEFGHIJK':put('16_Comparacion_Nacional_SIES',f'{cc}{r}',None)
# Build privacy-safe transport copies. Only original XML/style/objects are transported here.
z=zipfile.ZipFile(T);parts={n:z.read(n) for n in z.namelist()};ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};q=lambda s:'{'+ns['m']+'}'+s
ss=[]
if 'xl/sharedStrings.xml' in parts:
 st=E.fromstring(parts['xl/sharedStrings.xml']);ss=[''.join(el.itertext()) for el in st]
wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels']);ri={x.get('Id'):x.get('Target') for x in rels}
paths={}
for sheet in wb.find(q('sheets')):
 name=sheet.get('name');path=ri[sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')];path=path.lstrip('/') if path.startswith('/') else 'xl/'+path
 paths[name]=path
 if name in renames:sheet.set('name',renames[name])
for path in [n for n in parts if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')]:
 tree=E.fromstring(parts[path]);name=next(k for k,v in paths.items() if v==path)
 for row in tree.find(q('sheetData')):
  for c in list(row):
   addr=c.get('r');rn=int(re.search(r'\d+',addr).group())
   if (name==old9 and rn>=6) or (name==old11 and rn>=4):
    lim=N if name==old9 else N11
    if rn<=lim:
     for child in list(c):c.remove(child)
     c.attrib.pop('t',None)
    else:row.remove(c)
    continue
   if c.get('t')=='s':
    value=ss[int(c.find(q('v')).text)];c.remove(c.find(q('v')));c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=value
   ff=c.find(q('f'))
   if ff is not None and ff.text:
    for a,b in renames.items():ff.text=ff.text.replace(a,b)
 for el in tree.iter():
  if el.text:
   for a,b in renames.items():el.text=el.text.replace(a,b)
  for at,v in list(el.attrib.items()):
   for a,b in renames.items():v=v.replace(a,b)
   el.set(at,v)
 if name==s14:
  merges=tree.find(q('mergeCells'));E.SubElement(merges,q('mergeCell'),ref='L4:P4');merges.set('count',str(len(merges)))
  dvs=tree.find(q('dataValidations'))
  for dv in dvs:
   if dv.get('sqref')=='C4':dv.find(q('formula1')).text='$BL$2:$BL$4';dv.set('prompt','Seleccione una cohorte o ambas.')
  dv=E.SubElement(dvs,q('dataValidation'),type='list',sqref='L4',showErrorMessage='1',errorTitle='Seguimiento',error='Seleccione el seguimiento de la lista.');E.SubElement(dv,q('formula1')).text='$BK$2:$BK$5';dvs.set('count',str(len(dvs)))
 parts[path]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['xl/workbook.xml']=E.tostring(wb,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['xl/sharedStrings.xml']=E.tostring(E.Element(q('sst'),nsmap={None:ns['m']},count='0',uniqueCount='0'),xml_declaration=True,encoding='UTF-8',standalone=True)
# Remove stale calculation chain (its rows belonged to deleted personal input cells).
if 'xl/calcChain.xml' in parts:
 del parts['xl/calcChain.xml']
 for path in ['[Content_Types].xml','xl/_rels/workbook.xml.rels']:
  tree=E.fromstring(parts[path])
  for el in list(tree):
   if 'calcChain' in str(el.attrib):tree.remove(el)
  parts[path]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(P/'maqueta_sin_personales.xlsx','w',zipfile.ZIP_DEFLATED) as out:
 for n,b in parts.items():out.writestr(n,b)
# Read-only cached visual reference. It is never a deliverable and contains no student input rows.
refparts=dict(parts)
for path in paths.values():
 tree=E.fromstring(refparts[path])
 for c in tree.iter(q('c')):
  ff=c.find(q('f'))
  if ff is not None:
   c.remove(ff)
   if c.get('t')=='str':
    vv=c.find(q('v'));val=vv.text if vv is not None else ''
    if vv is not None:c.remove(vv)
    c.set('t','inlineStr');E.SubElement(E.SubElement(c,q('is')),q('t')).text=val
 refparts[path]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(P/'referencia_visual.xlsx','w',zipfile.ZIP_DEFLATED) as out:
 for n,b in refparts.items():out.writestr(n,b)
config={'template':str(T),'scaffold':str(P/'maqueta_sin_personales.xlsx'),'out':str(O/'RETENCION_INTERMEDIA_IPSS_MAQUETA_T1_20261021_153436.xlsx'),'edits':edits,'renames':renames,'groups':len(d['groups']),'last9':N,'last11':N11,'control':d['control'],'allcodes':allcodes,'codes2025':codes2025}
(O/'soporte/config.json').write_text(json.dumps(config,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'groups':len(d['groups']),'codes_2025':len(codes2025),'all_cohort_codes':len(allcodes),'edits':{k:len(v) for k,v in edits.items()},'privacy_transport':True},ensure_ascii=False))
