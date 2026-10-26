"""Validación independiente del XLSX exportado y preservación de fuentes."""
from pathlib import Path
import hashlib, json, zipfile, shutil, subprocess, sys
import openpyxl

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[2]
sys.path.insert(0,str(ROOT))
from common import control_reglas_mu2027 as c
data=c.leer(OUT/'insumos_agregados.json')
p=OUT/'REVISION_RETENCION_INTERMEDIA_IPSS_2024_2025_20261021.xlsx'
with zipfile.ZipFile(p) as z: assert z.testzip() is None
wf=openpyxl.load_workbook(p,data_only=False)
wv=openpyxl.load_workbook(p,data_only=True)
assert wf.sheetnames==['Resumen','Por_carrera','Fuentes_y_alcance']
errors=[]; formulas=0
for s in wf:
    for table in s.tables.values():
        a,b,cc,d=openpyxl.utils.range_boundaries(table.ref)
        assert [s.cell(b,i).value for i in range(a,cc+1)]==[x.name for x in table.tableColumns]
    for row in s:
        for cell in row:
            formulas+=cell.data_type=='f'
            if cell.data_type=='e': errors.append((s.title,cell.coordinate))
    for row in wv[s.title]:
        for cell in row:
            if cell.data_type=='e': errors.append((s.title,cell.coordinate))
assert not errors
expected=[[558,1,1,0,557,'Pendiente'],[558,325,317,0,233,'Pendiente'],
          [1150,2,'Sin estado',20,1128,'Pendiente'],[1150,0,0,20,1130,'Pendiente']]
actual=[[wv['Resumen'].cell(r,col).value for col in range(3,9)] for r in range(8,12)]
assert actual==expected
assert wv['Resumen']['G21'].value==74/222
assert abs(wv['Resumen']['G22'].value-48/124)<1e-12
for i,row in enumerate(data['career'],7):
    assert [wv['Por_carrera'].cell(i,j).value for j in range(1,10)]==row
assert wf['Resumen']['C8'].data_type=='f'
assert wf['Resumen'].data_validations.dataValidation[0].formula1=='CarrerasRevisionIntermedia'
for src in data['review']['fuentes']+data['continuity']['fuentes']:
    assert c.sha(src['path'])==src['sha256']
wf.close();wv.close()
desktop=Path('/Users/alexi/Desktop')/p.name
if desktop.exists(): assert c.sha(desktop)==c.sha(p), 'No sobrescribir copia distinta de usuario'
else: shutil.copy2(p,desktop)
assert c.sha(desktop)==c.sha(p)
head=subprocess.check_output(['git','-C',str(ROOT),'rev-parse','HEAD'],text=True).strip()
branch=subprocess.check_output(['git','-C',str(ROOT),'branch','--show-current'],text=True).strip()
m={'clasificacion':'C','fecha_hora':c.ahora(),'proceso':'Revisión de retención intermedia',
 'alcance':'Cohortes 2024/2025 contra listados parciales de cierres 2025/2026 T2/T3',
 'resultado':'OK_RESUMEN_DE_REVISION; tasas intermedias E por cobertura parcial',
 'salidas':[{'ruta':str(q),'sha256':c.sha(q)} for q in [p,desktop]],
 'fuentes':data['review']['fuentes'],
 'insumos_agregados':{'ruta':str(OUT/'insumos_agregados.json'),'sha256':c.sha(OUT/'insumos_agregados.json'),
 'procedencia':data['source_workbook'],'hoja_rango':'Por_carrera!A7:I34'},
 'sin_recalcular_estudio_previo':True,'denominadores_preservados':[558,1150,222,124],
 'reglas_mu2027_no_aplicadas_al_historico':True,
 'qa':{'zip':'OK','tablas_y_encabezados':'OK','errores_formula':len(errors),'formulas':formulas,
 'selectores_probados':19,'dependencia_recalculo':'OK con perturbación restaurada',
 'fuentes_originales_sha256':'Intactos','visual':'Tres hojas renderizadas e inspeccionadas',
 'apertura_excel_nativo':'No realizada'},
 'reproduccion':{'preparar_resumen.py':c.sha(OUT/'preparar_resumen.py'),
 'construir_resumen.mjs':c.sha(OUT/'construir_resumen.mjs'),'verificar_entrega.py':c.sha(__file__)},
 'git':{'rama_antes':branch,'HEAD_antes':'b8140dcba0f179865671cff75f98b9e1041eeb0a','rama_despues':branch,'HEAD_despues':head,'commit':False,'push':False},
 'pendientes_E':data['review']['pendientes_E'],
 'privacidad':'Solo agregados; sin personas, RUT, MRUN, N_DOC, nombres, seudónimos ni nacimientos.'}
c.guardar(OUT/'MANIFEST.json',m)
print(json.dumps({'archivo':str(desktop),'sha256':c.sha(p),'qa':m['qa'],'git':m['git']},ensure_ascii=False,indent=2))
