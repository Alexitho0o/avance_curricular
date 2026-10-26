"""Append only the Artifact-authored sheet; keep all old sheet parts byte-identical."""
import copy, datetime, hashlib, json, pathlib, shutil, subprocess, zipfile
from lxml import etree as E
import openpyxl

TMP=pathlib.Path('/tmp/ipss_escenario80_20261021_100204')
REPO=pathlib.Path('/Users/alexi/Documents/GitHub/avance_curricular')
BASE=REPO/'outputs/retencion_intermedia_t2_t3/20261021_085902/RETENCION_INTERMEDIA_IPSS_PRESENCIAS_Y_PENDIENTES_20261021_085902.xlsx'
STAMP=datetime.datetime.now().astimezone().strftime('%Y%m%d_%H%M%S')
OUT=REPO/'outputs/retencion_intermedia_t2_t3'/STAMP
OUT.mkdir()
FINAL=OUT/f'RETENCION_INTERMEDIA_IPSS_ESCENARIO_80_{STAMP}.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
def q(x):return f'{{{NS}}}{x}'
def sha(p):return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def xml(x):return E.fromstring(x)
def ser(x):return E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)

assert sha(BASE)=='340f9a2126973e0ab3a11172eaa525935ab5991c8ad10dcdb321cc721d0b8e69'
with zipfile.ZipFile(BASE) as z,zipfile.ZipFile(TMP/'candidate.xlsx') as c:
    original={n:z.read(n) for n in z.namelist()}
    styles=xml(original['xl/styles.xml'])
    cs=xml(c.read('xl/styles.xml'))
    nm={}; dest=styles.find(q('numFmts')); src=cs.find(q('numFmts'))
    if dest is None:
        dest=E.Element(q('numFmts'),count='0');styles.insert(0,dest)
    used={int(t.get('numFmtId')) for t in dest}
    for el in src if src is not None else []:
        old=int(el.get('numFmtId'));new=max({163}|used)+1;used.add(new);nm[old]=new
        cp=copy.deepcopy(el);cp.set('numFmtId',str(new));dest.append(cp)
    dest.set('count',str(len(dest)))
    offsets={}
    for kind in ['fonts','fills','borders','cellStyleXfs','cellXfs','dxfs']:
        dst=styles.find(q(kind)); source=cs.find(q(kind))
        if dst is None:
            dst=E.SubElement(styles,q(kind),count='0')
        offsets[kind]=len(dst)
        for el in source if source is not None else []:
            cp=copy.deepcopy(el)
            if kind in ('cellStyleXfs','cellXfs'):
                for attr,section in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
                    if attr in cp.attrib:cp.set(attr,str(int(cp.get(attr))+offsets[section]))
                if 'numFmtId' in cp.attrib:
                    n=int(cp.get('numFmtId'));cp.set('numFmtId',str(nm.get(n,n)))
                if kind=='cellXfs' and 'xfId' in cp.attrib:cp.set('xfId',str(int(cp.get('xfId'))+offsets['cellStyleXfs']))
            for fmt in cp.iter(q('numFmt')):
                n=int(fmt.get('numFmtId'));fmt.set('numFmtId',str(nm.get(n,n)))
            dst.append(cp)
        dst.set('count',str(len(dst)))
    # Match the new worksheet through its relationship, never assume exported numbering.
    cw=xml(c.read('xl/workbook.xml'));rels=xml(c.read('xl/_rels/workbook.xml.rels'))
    sheet=next(t for t in cw.find(q('sheets')) if t.get('name')=='Escenario_80')
    rid=sheet.get(f'{{{RNS}}}id')
    target=next(t.get('Target') for t in rels if t.get('Id')==rid)
    sheetpath=target.lstrip('/') if target.startswith('/') else 'xl/'+target
    sn=xml(c.read(sheetpath))
    assert not any(el.tag==q('drawing') for el in sn.iter())
    ss=xml(c.read('xl/sharedStrings.xml')) if 'xl/sharedStrings.xml' in c.namelist() else []
    for el in sn.iter():
        if el.tag==q('c'):
            if 's' in el.attrib:el.set('s',str(int(el.get('s'))+offsets['cellXfs']))
            if el.get('t')=='s':
                idx=int(el.find(q('v')).text);el.remove(el.find(q('v')))
                el.set('t','inlineStr');inline=E.SubElement(el,q('is'))
                for child in ss[idx]:inline.append(copy.deepcopy(child))
        elif el.tag==q('col') and 'style' in el.attrib:el.set('style',str(int(el.get('style'))+offsets['cellXfs']))
        elif el.tag==q('row') and 's' in el.attrib:el.set('s',str(int(el.get('s'))+offsets['cellXfs']))
        elif el.tag==q('cfRule') and 'dxfId' in el.attrib:el.set('dxfId',str(int(el.get('dxfId'))+offsets['dxfs']))
    wb=xml(original['xl/workbook.xml']);wr=xml(original['xl/_rels/workbook.xml.rels']);ct=xml(original['[Content_Types].xml'])
    ids={el.get('Id') for el in wr};newid=1
    while f'rId{newid}' in ids:newid+=1
    newrid=f'rId{newid}'; sheets=wb.find(q('sheets'));sid=max(int(s.get('sheetId')) for s in sheets)+1
    newpath=f'xl/worksheets/sheet{sid}.xml';assert newpath not in original
    E.SubElement(sheets,q('sheet'),name='Escenario_80',sheetId=str(sid),attrib={f'{{{RNS}}}id':newrid})
    E.SubElement(wr,f'{{{PKG}}}Relationship',Id=newrid,Type=RNS+'/worksheet',Target=f'worksheets/sheet{sid}.xml')
    E.SubElement(ct,f'{{{CT}}}Override',PartName='/'+newpath,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    replacements={'xl/styles.xml':ser(styles),'xl/workbook.xml':ser(wb),'xl/_rels/workbook.xml.rels':ser(wr),'[Content_Types].xml':ser(ct),newpath:ser(sn)}
    with zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as out:
        for n,content in original.items():out.writestr(n,replacements.get(n,content))
        out.writestr(newpath,replacements[newpath])

with zipfile.ZipFile(FINAL) as z:
    assert z.testzip() is None
    preserved=[n for n in original if n not in replacements]
    assert all(z.read(n)==original[n] for n in preserved)
    for n in z.namelist():
        if n.endswith('.xml'):xml(z.read(n))
# Read-only independent examination, not openpyxl authoring.
a=openpyxl.load_workbook(BASE,data_only=False)
b=openpyxl.load_workbook(FINAL,data_only=False)
assert b.sheetnames==a.sheetnames+['Escenario_80']
for name in a.sheetnames:
    sa,sb=a[name],b[name]
    assert sa.max_row==sb.max_row and sa.max_column==sb.max_column
    for row in sa:
        for ca in row:
            cb=sb[ca.coordinate]
            assert ca.value==cb.value,(name,ca.coordinate,ca.value,cb.value)
            assert ca._style==cb._style,(name,ca.coordinate,'style')
    assert str(sa.merged_cells)==str(sb.merged_cells)
    assert list(sa.tables)==list(sb.tables)
cached=openpyxl.load_workbook(FINAL,data_only=True)
sn=cached['Escenario_80']
expect={'B8':558,'C8':326,'E8':560,'F8':328,'D8':326/558,'G8':328/560,'B9':1150,'C9':668,'E9':1123,'F9':668,'D9':668/1150,'G9':668/1123,'E28':80,'E32':80,'E52':80,'F52':53,'G52':11,'H52':16,'F19':0,'F20':0}
expect.update({f'E{r}':0 for r in range(55,60)})
for cell,val in expect.items():
    assert isinstance(sn[cell].value,(int,float)) and abs(sn[cell].value-val)<1e-12,(cell,sn[cell].value,val)
errors=[(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type=='e']
assert not errors,errors
assert sha(BASE)=='340f9a2126973e0ab3a11172eaa525935ab5991c8ad10dcdb321cc721d0b8e69'
desktop=pathlib.Path('/Users/alexi/Desktop')/FINAL.name
assert not desktop.exists();shutil.copy2(FINAL,desktop);assert sha(FINAL)==sha(desktop)
for name in ['build.mjs','preserve_and_verify.py']:
    shutil.copy2(TMP/name,OUT/name)
head=subprocess.check_output(['git','-C',str(REPO),'rev-parse','HEAD'],text=True).strip()
assert head=='80ca13a5a249a1a2bd53a699f5f73689b1f53040'
manifest={'hito':'H001-RETENCION-IPSS162','fecha':datetime.datetime.now().astimezone().isoformat(),'clasificacion':'B observaciones agregadas; C fórmulas; D preservación de actuales; E simulación no aprobada','entrada':{'ruta':str(BASE),'sha256':sha(BASE)},'escenario_entrada':{'ruta':str(REPO/'outputs/retencion_intermedia_t2_t3/20261021_095123_escenario_80/RESULTADOS_ESCENARIO_AGREGADOS.json'),'sha256':'4a4f0f7183d4a1bcff507085399920cbaa3e37c5829bf2d11ced58ef38d0681c'},'salida':{'ruta':str(FINAL),'sha256':sha(FINAL)},'copia_escritorio':{'ruta':str(desktop),'sha256':sha(desktop)},'hoja_agregada':'Escenario_80','originales_sin_modificacion':True,'denominadores_actuales':[558,1150,222,124],'denominadores_simulados_solo_hoja_nueva':[560,1123],'QA':{'resultado':'OK','zip_integro':True,'xml_bien_formado':True,'hojas_previas_identicas':a.sheetnames,'partes_originales_byte_identicas':preserved,'errores_formula_en_caches':errors,'recalculo':'Artifact Tool; controles numéricos independientes; sensibilidad E36/E33 con restauración antes de exportar','validacion_visual':'Resumen, entradas/desglose y pendientes/fuentes revisados','native_excel':'No se realizó apertura automatizada en Excel'},'rama':'feature/ire-2026','HEAD_antes':head,'HEAD_despues':head,'commit':False,'push':False,'datos_individuales_agregados':False,'pendientes':['Fundamento de la asignación de origen 2025 a ingresos 2024','Semestre de origen correcto','78 personas sin fila MU2024','Cobertura/fecha/estados de cierres T2/T3'], 'scripts':{name:sha(OUT/name) for name in ['build.mjs','preserve_and_verify.py']}}
(OUT/'MANIFEST_ESCENARIO_80_EXCEL.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
(TMP/'final_paths.json').write_text(json.dumps({'excel':str(FINAL),'desktop':str(desktop),'manifest':str(OUT/'MANIFEST_ESCENARIO_80_EXCEL.json'),'sha256':sha(FINAL)},ensure_ascii=False,indent=2))
print(json.dumps({'excel':str(FINAL),'desktop':str(desktop),'sha256':sha(FINAL),'QA':'OK','hojas':b.sheetnames,'errores_formula':0,'originales_preservados':len(preserved)},ensure_ascii=False,indent=2))
