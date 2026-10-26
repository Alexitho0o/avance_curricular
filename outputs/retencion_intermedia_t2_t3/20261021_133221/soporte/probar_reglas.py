"""Read-only rule audit; synthetic probes and aggregate evidence only."""
import ast, hashlib, json, sys, unittest
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
ROOT = Path('/Users/alexi/Documents/GitHub/avance_curricular')
sys.path.insert(0, str(ROOT/'procesos/caracterizacion'))
from caracterizacion_retencion import calcular_retencion_longitudinal

source=ROOT/'procesos/caracterizacion/tests/test_hito4b_sectorizado_sies_retencion.py'
tree=ast.parse(source.read_text())
cls=next(n for n in tree.body if isinstance(n,ast.ClassDef) and n.name=='Hito4BRetencionTests')
ns={'pd':pd, 'unittest':unittest, 'calcular_retencion_longitudinal':calcular_retencion_longitudinal}
exec(compile(ast.Module(body=[cls],type_ignores=[]),str(source),'exec'),ns)
result=unittest.TestResult()
unittest.defaultTestLoader.loadTestsFromTestCase(ns['Hito4BRetencionTests']).run(result)
probes=[]
for base,dest in [('TITULADO',None),('VIGENTE','TITULADO'),('VIGENTE','EGRESADO'),('VIGENTE',None),('VIGENTE','SUSPENDIDO')]:
    rows=[dict(RUT='12.345.678-5',CODCLI='SYNTHETIC',CODCARPR='TEST',ANIO=2025,PERIODO=2,ESTADO=base)]
    if dest: rows.append(dict(rows[0],ANIO=2026,ESTADO=dest))
    z=calcular_retencion_longitudinal(pd.DataFrame(rows),anio_base=2025,anio_destino=2026,periodo_base=2,periodo_destino=2,fecha_corte='2026-07-15')
    d=z.detalle.iloc[0]
    probes.append({'estado_base_sintetico':base,'estado_destino_sintetico':dest,'resultado':d.RESULTADO_RETENCION,'elegible':bool(d.ES_ELEGIBLE)})
p=ROOT/'procesos/caracterizacion/gobernanza/sies/RETENCION/Retencion_ipss_2022_2024_GRANULAR.xlsx'
wb=load_workbook(p,read_only=True,data_only=False)
print('WORKBOOK',hashlib.sha256(p.read_bytes()).hexdigest(),wb.sheetnames)
for s in wb:
    if 'dic' in s.title.lower():
        for i,row in enumerate(s.iter_rows(min_row=1,max_row=74,max_col=1,values_only=True),1):
            if row[0] is not None: print(f'{s.title}!A{i}',row[0])
wb.close()
out={'existing_retention_test': {'testsRun':result.testsRun,'errors':len(result.errors),'failures':len(result.failures),'method':'AST extraction of unchanged Hito4BRetencionTests, isolates retention from unrelated missing unidecode dependency'},'synthetic_probes':probes,'full_suite_limitation':'Full test module import blocked by missing unidecode in bundled runtime; isolated retention test completed.'}
Path('/tmp/retencion_rules_probe.json').write_text(json.dumps(out,ensure_ascii=False,indent=2))
print(json.dumps(out,ensure_ascii=False,indent=2))
