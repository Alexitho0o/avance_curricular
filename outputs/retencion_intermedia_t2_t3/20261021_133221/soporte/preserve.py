"""Transport Artifact-authored edits while preserving original XLSX features.

No individual data are handled. openpyxl is read-only for independent QA.
"""
import copy, hashlib, json, pathlib, zipfile
from lxml import etree as E
import openpyxl

TMP=pathlib.Path('/tmp/retencion_revision_20261021')
cfg=json.loads((TMP/'config.json').read_text())
BASE=pathlib.Path(cfg['base']); FINAL=pathlib.Path(cfg['output'])
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PKG='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
def q(x):return '{'+NS+'}'+x
def sha(p):return hashlib.sha256(pathlib.Path(p).read_bytes()).hexdigest()
def xml(x):return E.fromstring(x)
def ser(x):return E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)
def sheetpaths(parts):
    wb=xml(parts['xl/workbook.xml']);rel=xml(parts['xl/_rels/workbook.xml.rels'])
    targets={e.get('Id'):e.get('Target') for e in rel}
    return {e.get('name'):(lambda t:t.lstrip('/') if t.startswith('/') else 'xl/'+t)(targets[e.get('{'+RNS+'}id')]) for e in wb.find(q('sheets'))}
assert sha(BASE)=='12bf04615558554a30275b4ee8c6197617e7195c059c4c87277a91c4697e719d'
with zipfile.ZipFile(BASE) as z,zipfile.ZipFile(TMP/'candidate.xlsx') as c:
    original={n:z.read(n) for n in z.namelist()}; candidate={n:c.read(n) for n in c.namelist()}
styles=xml(original['xl/styles.xml']);cs=xml(candidate['xl/styles.xml'])
nm={};dest=styles.find(q('numFmts'));src=cs.find(q('numFmts'))
if dest is None:dest=E.Element(q('numFmts'),count='0');styles.insert(0,dest)
used={int(t.get('numFmtId')) for t in dest}
for el in src if src is not None else []:
    old=int(el.get('numFmtId'));new=max({163}|used)+1;used.add(new);nm[old]=new
    cp=copy.deepcopy(el);cp.set('numFmtId',str(new));dest.append(cp)
dest.set('count',str(len(dest)))
offsets={}
for kind in ['fonts','fills','borders','cellStyleXfs','cellXfs','dxfs']:
    dst=styles.find(q(kind));source=cs.find(q(kind))
    if dst is None:dst=E.SubElement(styles,q(kind),count='0')
    offsets[kind]=len(dst)
    for el in source if source is not None else []:
        cp=copy.deepcopy(el)
        if kind in ('cellStyleXfs','cellXfs'):
            for attr,section in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:
                if attr in cp.attrib:cp.set(attr,str(int(cp.get(attr))+offsets[section]))
            if 'numFmtId' in cp.attrib:
                n=int(cp.get('numFmtId'));cp.set('numFmtId',str(nm.get(n,n)))
            if kind=='cellXfs' and 'xfId' in cp.attrib:cp.set('xfId',str(int(cp.get('xfId'))+offsets['cellStyleXfs']))
        for fmt in cp.iter(q('numFmt')):
            n=int(fmt.get('numFmtId'));fmt.set('numFmtId',str(nm.get(n,n)))
        dst.append(cp)
    dst.set('count',str(len(dst)))
ss=xml(candidate['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in candidate else []
def inline(cell):
    if cell.get('t')=='s':
        idx=int(cell.find(q('v')).text);cell.remove(cell.find(q('v')));cell.set('t','inlineStr')
        target=E.SubElement(cell,q('is'))
        for child in ss[idx]:target.append(copy.deepcopy(child))
    return cell
op=sheetpaths(original);cp=sheetpaths(candidate)
replacements={'xl/styles.xml':ser(styles)};cache_changes=[]
for name,path in op.items():
    tree=xml(original[path]);ctree=xml(candidate[cp[name]])
    cells={el.get('r'):el for el in tree.iter(q('c'))};ccells={el.get('r'):el for el in ctree.iter(q('c'))}
    changed=False
    for address in cfg['edits'].get(name,{}):
        old=cells[address];new=inline(copy.deepcopy(ccells[address]));new.set('s',old.get('s','0'))
        old.getparent().replace(old,new);cells[address]=new;changed=True
    # All original formula expressions remain literal. Refresh caches only when changed.
    for address,old in cells.items():
        if old.find(q('f')) is None or address in cfg['edits'].get(name,{}):continue
        new=ccells.get(address)
        if new is None:raise AssertionError((name,address,'candidate formula absent'))
        ov=old.find(q('v'));nv=new.find(q('v'))
        old_value=(old.get('t'),None if ov is None else ov.text)
        new_value=(new.get('t'),None if nv is None else nv.text)
        if old_value!=new_value:
            if ov is not None:old.remove(ov)
            if nv is not None:old.append(copy.deepcopy(nv))
            if new.get('t') is None:old.attrib.pop('t',None)
            else:old.set('t',new.get('t'))
            cache_changes.append([name,address,old_value,new_value]);changed=True
    if changed:replacements[path]=ser(tree)
newname='Revision_agregada';sn=xml(candidate[cp[newname]])
assert sn.find(q('drawing')) is None and sn.find(q('tableParts')) is None
for el in sn.iter():
    if el.tag==q('c'):
        if 's' in el.attrib:el.set('s',str(int(el.get('s'))+offsets['cellXfs']))
        inline(el)
    elif el.tag==q('col') and 'style' in el.attrib:el.set('style',str(int(el.get('style'))+offsets['cellXfs']))
    elif el.tag==q('row') and 's' in el.attrib:el.set('s',str(int(el.get('s'))+offsets['cellXfs']))
    elif el.tag==q('cfRule') and 'dxfId' in el.attrib:el.set('dxfId',str(int(el.get('dxfId'))+offsets['dxfs']))
wb=xml(original['xl/workbook.xml']);wr=xml(original['xl/_rels/workbook.xml.rels']);ct=xml(original['[Content_Types].xml'])
ids={el.get('Id') for el in wr};rid=1
while f'rId{rid}' in ids:rid+=1
sheets=wb.find(q('sheets'));sid=max(int(s.get('sheetId')) for s in sheets)+1
newpath=f'xl/worksheets/sheet{sid}.xml'
E.SubElement(sheets,q('sheet'),name=newname,sheetId=str(sid),attrib={'{'+RNS+'}id':f'rId{rid}'})
E.SubElement(wr,'{'+PKG+'}Relationship',Id=f'rId{rid}',Type=RNS+'/worksheet',Target=f'worksheets/sheet{sid}.xml')
E.SubElement(ct,'{'+CT+'}Override',PartName='/'+newpath,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
replacements.update({'xl/workbook.xml':ser(wb),'xl/_rels/workbook.xml.rels':ser(wr),'[Content_Types].xml':ser(ct),newpath:ser(sn)})
FINAL.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(FINAL,'w',zipfile.ZIP_DEFLATED) as z:
    for n,content in original.items():z.writestr(n,replacements.get(n,content))
    z.writestr(newpath,replacements[newpath])
a=openpyxl.load_workbook(BASE,data_only=False);b=openpyxl.load_workbook(FINAL,data_only=False)
assert b.sheetnames==a.sheetnames+[newname]
for name in a.sheetnames:
    sa,sb=a[name],b[name]
    assert (sa.max_row,sa.max_column)==(sb.max_row,sb.max_column)
    assert str(sa.merged_cells)==str(sb.merged_cells)
    assert list(sa.tables)==list(sb.tables)
    assert str(sa.data_validations)==str(sb.data_validations)
    assert sa.freeze_panes==sb.freeze_panes
    for row in sa:
        for ca in row:
            cb=sb[ca.coordinate]
            if ca.coordinate not in cfg['edits'].get(name,{}):assert ca.value==cb.value,(name,ca.coordinate,'unapproved value change')
            # openpyxl synthesizes merged-cell border IDs; compare actual styles.
            for attr in ['font','fill','border','alignment','protection','number_format']:
                assert copy.copy(getattr(ca,attr))==copy.copy(getattr(cb,attr)),(name,ca.coordinate,attr,'style changed')
with zipfile.ZipFile(FINAL) as z:
    assert z.testzip() is None
    preserved=[n for n in original if n not in replacements]
    assert all(z.read(n)==original[n] for n in preserved)
    assert z.read(op['Escenario_80'])==original[op['Escenario_80']]
    original_styles=xml(original['xl/styles.xml']);final_styles=xml(z.read('xl/styles.xml'))
    for section in ['fonts','fills','borders','cellStyleXfs','cellXfs','dxfs','numFmts']:
        old_section=original_styles.find(q(section));new_section=final_styles.find(q(section))
        if old_section is not None:
            assert all(E.tostring(u)==E.tostring(v) for u,v in zip(old_section,new_section)),('original_style_definition',section)
    for name in z.namelist():
        if name.endswith('.xml'):xml(z.read(name))
cached=openpyxl.load_workbook(FINAL,data_only=True)
errors=[(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type=='e']
assert not errors,errors
for sheet,entries in cfg['expected'].items():
    for address,value in entries.items():
        actual=cached[sheet][address].value
        if isinstance(value,(int,float)):assert isinstance(actual,(int,float)) and abs(actual-value)<1e-11,(sheet,address,actual,value)
        else:assert actual==value,(sheet,address,actual,value)
assert sha(BASE)=='12bf04615558554a30275b4ee8c6197617e7195c059c4c87277a91c4697e719d'
qa={'zip_xml':'OK','original_sha256':sha(BASE),'output_sha256':sha(FINAL),'preserved_package_parts':preserved,'cache_changes':cache_changes,'approved_edits':cfg['edits'],'formula_errors':errors,'all_original_styles_tables_filters_validations_preserved':True,'escenario80_byte_identical':True,'denominadores':[558,1150,222,124],'all_expected_values_pass':True}
(TMP/'qa_package.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2))
print(json.dumps({'file':str(FINAL),'sha256':sha(FINAL),'checks':'OK','cache_changes':len(cache_changes),'sheets':b.sheetnames},ensure_ascii=False))
