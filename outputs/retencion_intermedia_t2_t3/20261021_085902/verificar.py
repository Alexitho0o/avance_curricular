from pathlib import Path
import hashlib,json,zipfile,shutil,subprocess
import openpyxl
p=Path(__file__).resolve().parent
root=p.parents[2]
e=json.loads((p/'evidencia_actualizacion.json').read_text())
src=Path(e['origen']);out=Path(e['salida'])
sha=lambda f:hashlib.sha256(Path(f).read_bytes()).hexdigest()
assert sha(src)=='1fed8a17c0b6977c579b14a86c012b7160170648b46e638f9009f33b0c0c85a0'
with zipfile.ZipFile(out) as z:assert z.testzip() is None
a=openpyxl.load_workbook(src,data_only=False);b=openpyxl.load_workbook(out,data_only=False);v=openpyxl.load_workbook(out,data_only=True)
allowed={'Resumen':{'A5','A13','A14','A27','A29','A30','A31',*[f'I{i}' for i in range(8,12)]},'Por_carrera':{'A4','A39'},'Fuentes_y_alcance':{'A20','A22','A25'}}
changes=[];nformula=0
for s in a:
 for row in s:
  for c in row:
   other=b[s.title][c.coordinate]
   if c.value!=other.value:
    assert c.coordinate in allowed[s.title] or (s.title=='Fuentes_y_alcance' and 34<=c.row<=41 and c.column<=6),(s.title,c.coordinate,c.value,other.value)
    changes.append([s.title,c.coordinate])
   if c.data_type=='f':assert c.value==other.value;nformula+=1
 assert s.freeze_panes==b[s.title].freeze_panes
 assert set(s.tables)==set(b[s.title].tables)
for s in b:
 for t in s.tables.values():
  lo,rr,hi,end=openpyxl.utils.range_boundaries(t.ref)
  assert [s.cell(rr,j).value for j in range(lo,hi+1)]==[c.name for c in t.tableColumns]
for s in v:
 for row in s:
  for c in row:assert c.data_type!='e',(s.title,c.coordinate,c.value)
assert b.sheetnames==a.sheetnames+['Bases_por_confirmar']
assert [[v['Resumen'].cell(r,c).value for c in range(3,9)] for r in range(8,12)]==e['controles']
assert [v['Resumen'][x].value for x in ['C8','C10','B21','B22']]==[558,1150,222,124]
assert b['Resumen'].data_validations.dataValidation[0].formula1==a['Resumen'].data_validations.dataValidation[0].formula1
for x in e['fuentes']:assert sha(x[4])==x[3]
for s in [a,b,v]:s.close()
desk=Path('/Users/alexi/Desktop')/out.name
assert not desk.exists() or sha(desk)==sha(out)
shutil.copy2(out,desk);assert sha(desk)==sha(out)
git=lambda *args:subprocess.check_output(['git','-C',str(root),*args],text=True).strip()
manifest={'clasificacion':'C','fecha':e['fecha'],'entrada':{'ruta':str(src),'sha256':sha(src)},'salidas':[{'ruta':str(f),'sha256':sha(f)} for f in [out,desk]],'fuentes':e['fuentes'],'cambios':changes,'denominadores_preservados':e['denominadores'],'presencias_conservadas':True,'tasas_intermedias':'E: pendientes por cobertura o fecha no acreditadas','QA':{'zip':'OK','formulas_originales_preservadas':nformula,'errores_formula':0,'selector_carrera':'19 selecciones contrastadas con detalle','visual':'Resumen y hoja nueva renderizados y revisados','excel_nativo':'No probado'},'git':{'rama_antes':'feature/ire-2026','HEAD_antes':'80ca13a5a249a1a2bd53a699f5f73689b1f53040','rama_despues':git('branch','--show-current'),'HEAD_despues':git('rev-parse','HEAD'),'commit':False,'push':False},'privacidad':'Solo agregados. Sin identificadores individuales.','implementacion':{'archivo':'actualizar.mjs','sha256':sha(p/'actualizar.mjs')}}
(p/'MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps({'salidas':manifest['salidas'],'QA':manifest['QA'],'git':manifest['git']},ensure_ascii=False,indent=2))
