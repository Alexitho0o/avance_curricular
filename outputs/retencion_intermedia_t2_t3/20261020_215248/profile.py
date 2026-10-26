from pathlib import Path
from collections import Counter, defaultdict
import json, hashlib, re, datetime
import openpyxl

ROOT=Path('/Users/alexi/Documents/GitHub/avance_curricular')
TMP=Path(__file__).resolve().parent
rawpath=ROOT/'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
engpath=ROOT/'outputs/retencion_primer_anio_ipss_162/20260930_000000/RETENCION_PRIMER_ANIO_IPSS_162_DIAGNOSTICO_BRECHAS_20260930.xlsx'
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def txt(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip().upper()
def key(n,d=None):
    v=txt(n).replace('.','').replace(' ','')
    if d is None:
        if '-' not in v:return None
        v,d=v.rsplit('-',1)
    else:d=txt(d)
    if not re.fullmatch(r'\d+',v) or not re.fullmatch(r'[0-9K]',d):return None
    return str(int(v))+'-'+d
def dvok(k):
    if not k:return False
    n,d=k.split('-'); r=11-sum(int(c)*(2+i%6) for i,c in enumerate(n[::-1]))%11
    return d==('0' if r==11 else 'K' if r==10 else str(r))
def records(s):
    rows=list(s.values); h=[str(v).strip() for v in rows[0]]
    return [dict(zip(h,r),_row=i) for i,r in enumerate(rows[1:],2) if any(v is not None for v in r)]
rawbook=openpyxl.load_workbook(rawpath,read_only=True,data_only=True)
raw=records(rawbook['BASE_RETENCION_MU']); rawby={r['_row']:r for r in raw}; rawbook.close()
eb=openpyxl.load_workbook(engpath,read_only=True,data_only=True)
engine=records(eb['02_COHORTE_ELEGIBILIDAD'])
catalog=records(eb['01_CATALOGO_CARRERA'])
catidx=defaultdict(list)
for c in catalog:catidx[str(c['ANO'])+'|'+str(c['CODIGO_UNICO'])].append(c)
by_person=defaultdict(list)
for r in raw:by_person[(r['ANIO_INFORMADO'],r['TIPO_DOC'],r['N_DOC'],r['DV'])].append(r)
# The original engine has formulas but no cached values. Evaluate its observed logic
# without saving or altering it; keep the exact filters and identify this as C.
engine=[]
for r in raw:
    y=r['ANIO_INFORMADO']
    if y not in (2024,2025) or r['ANIO_ING_ORI']!=y or r['SEM_ING_ORI'] not in (1,2):continue
    code=f"I162S{r['COD_SED']}C{r['COD_CAR']}J{r['JOR']}V{r['VERSION']}"
    cc=catidx['OFE_'+str(y)+'|'+code]; c=cc[0] if len(cc)==1 else {}
    mins={'Técnico de Nivel Superior':4,'Profesional Sin Licenciatura':6,'Profesional Con Licenciatura':8,'Licenciatura no conducente a título':8}.get(c.get('NIVEL_CARRERA'))
    dur=float(c['DURACION_ESTUDIOS']) if c.get('DURACION_ESTUDIOS') not in (None,'') else None
    status='VIGENCIA_0' if r['VIG']==0 else 'CATALOGO_NO_UNIVOCO' if len(cc)!=1 else 'OTRO_NIVEL' if c['NIVEL_GLOBAL']!='Pregrado' else 'PLAN_NO_REGULAR' if c['TIPO_CARRERA']!='Plan Regular' else 'PLAN_ESPECIAL' if c['PLAN_ESPECIAL']!='NO APLICA' else 'CATEGORIA_PENDIENTE' if mins is None else 'DURACION_INSUFICIENTE' if dur is None or dur<mins else 'ELEGIBLE'
    ident=(r['TIPO_DOC'],r['N_DOC'],r['DV'])
    same=lambda q:q['ANIO_ING_ORI']==r['ANIO_ING_ORI'] and q['SEM_ING_ORI']==r['SEM_ING_ORI']
    dup=sum(same(q) and (q['VIG'] or 0)>=1 for q in by_person[(y,*ident)]) if status=='ELEGIBLE' else 0
    nxt=[q for q in by_person[(y+1,*ident)] if (q['VIG'] or 0)>=1] if status=='ELEGIBLE' else []
    nco=sum(same(q) for q in nxt)
    final=status if status!='ELEGIBLE' else 'DUPLICADO_PERSONA_COHORTE' if dup>1 else 'COHORTE_INCONSISTENTE' if len(nxt)>nco else 'INCLUIDO'
    engine.append({'FILA_RAW':r['_row'],'ANIO_INFORMADO':y,'COHORTE_CANDIDATA':1,'CODIGO_UNICO':code,'NOMBRE_CARRERA':c.get('NOMBRE_CARRERA','Sin catálogo unívoco'),'AREA_DEL_CONOCIMIENTO':c.get('AREA_DEL_CONOCIMIENTO','Sin catálogo unívoco'),'ESTADO_FINAL':final,'N_CASOS':int(final=='INCLUIDO'),'ESTADO_ELEGIBILIDAD':status,'N_CATALOGO':len(cc),'NIVEL_GLOBAL':c.get('NIVEL_GLOBAL'),'TIPO_CARRERA':c.get('TIPO_CARRERA'),'PLAN_ESPECIAL':c.get('PLAN_ESPECIAL'),'NIVEL_CARRERA':c.get('NIVEL_CARRERA'),'DURACION_ESTUDIOS':dur,'MIN_SEMESTRES':mins,'DUPLICADOS_ACTIVOS':dup,'SIGUIENTE_ANY':len(nxt),'SIGUIENTE_MISMA_COHORTE':nco})
base=[]
for e in engine:
    if e['ANIO_INFORMADO'] not in (2024,2025) or e['COHORTE_CANDIDATA']!=1:continue
    r=rawby[e['FILA_RAW']]
    k=key(r['N_DOC'],r['DV']) if r['TIPO_DOC']=='R' else None
    base.append({'key':k,'dv_ok':dvok(k),'raw':r,'engine':e})
eb.close()
specs=[('T2_2025','*T2_2025*','T2_2025',2024,2025,2),('T3_2025','*T3_2025*','T3_2025',2024,2025,3),('T2_2026','*13_07*','Hoja1',2025,2026,2),('T3_2026','*02_10*','Base',2025,2026,3)]
sources=[]; closure=[]; inventory=[]
for sid,pat,selected,cohort,year,term in specs:
    p=next(Path('/Users/alexi/Downloads/cierre').glob(pat));w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    for s in w:
        rr=list(s.values); hh=[str(v).strip() for v in rr[0]]
        detail='RUT' in hh
        inventory.append({'source':sid,'file':p.name,'sheet':s.title,'rows_used':sum(any(v is not None for v in row) for row in rr[1:]),'dimensions':[s.max_row,s.max_column],'headers':hh,'role':'Detalle seleccionado' if s.title==selected else 'Otra hoja de detalle; no concatenada' if detail else 'Resumen/listado auxiliar; no usado como matrícula','selected':s.title==selected})
    rr=records(w[selected]); idx=defaultdict(list)
    for r in rr:
        k=key(r.get('RUT'));idx[k].append(r)
        closure.append({'source':sid,'sheet':selected,'cohort':cohort,'year':year,'term':term,'key':k,'dv_ok':dvok(k),'row':r['_row'],'entry_year':r.get('ANOINGRESO'),'entry_term':r.get('PERIODOINGRESO'),'state':r.get('ESTADOACADEMICO','NO INFORMADO'),'career':r.get('NOMBRE_L'),'career_code':r.get('CODCARPR',r.get('CODIGOCARRERA')),'enrol_year':r.get('ANOMATRICULA'),'enrol_term':r.get('PERIODOMATRICULA')})
    targets=[b for b in base if b['engine']['ANIO_INFORMADO']==cohort]
    inc=[b for b in targets if b['engine']['ESTADO_FINAL']=='INCLUIDO']
    match=[b for b in inc if b['key'] and b['key'] in idx]
    vig=[b for b in match if any(r.get('ESTADOACADEMICO')=='VIGENTE' for r in idx[b['key']])]
    sources.append({'id':sid,'path':str(p),'sha256':digest(p),'bytes':p.stat().st_size,'mtime':datetime.datetime.fromtimestamp(p.stat().st_mtime).astimezone().isoformat(),'sheet':selected,'cohort':cohort,'year':year,'term':term,'records':len(rr),'unique_keys':len([k for k in idx if k]),'invalid_keys':sum(key(r.get('RUT')) is None for r in rr),'invalid_dv':sum(not dvok(key(r.get('RUT'))) for r in rr),'entry_years':dict(Counter(txt(r.get('ANOINGRESO')) for r in rr)),'entry_terms':dict(Counter(txt(r.get('PERIODOINGRESO')) for r in rr)),'states':dict(Counter(txt(r.get('ESTADOACADEMICO','NO INFORMADO')) for r in rr)),'targets':len(targets),'included':len(inc),'match':len(match),'match_vig':len(vig),'match_invalid_dv':sum(not b['dv_ok'] for b in match)})
    w.close()
keys=sorted(set(b['key'] for b in base if b['key'])|set(c['key'] for c in closure if c['key']))
pseudo={k:f'P{i:05d}' for i,k in enumerate(keys,1)}
# Only pseudonyms and original source row references leave memory; no individual IDs in JSON/output.
baseout=[]
for b in base:
    r,e=b['raw'],b['engine']
    baseout.append({'person':pseudo.get(b['key'],f'SIN_LLAVE_{r["_row"]}'),'key_ok':bool(b['key']),'dv_ok':b['dv_ok'],'source_row':r['_row'],'year':r['ANIO_INFORMADO'],'origin_year':r['ANIO_ING_ORI'],'origin_term':r['SEM_ING_ORI'],'vig':r['VIG'],'unique_code':e['CODIGO_UNICO'],'career':e['NOMBRE_CARRERA'],'area':e['AREA_DEL_CONOCIMIENTO'],'annual_state':e['ESTADO_FINAL'],'annual_eligible':e['N_CASOS'],'eligibility_before_cohort':e['ESTADO_ELEGIBILIDAD'],'type_doc':r['TIPO_DOC'],'catalog_matches':e['N_CATALOGO'],'level':e['NIVEL_GLOBAL'],'plan':e['TIPO_CARRERA'],'special':e['PLAN_ESPECIAL'],'career_level':e['NIVEL_CARRERA'],'duration':e['DURACION_ESTUDIOS'],'min_semesters':e['MIN_SEMESTRES'],'duplicate_active':e['DUPLICADOS_ACTIVOS'],'annual_next_any':e['SIGUIENTE_ANY'],'annual_next_cohort':e['SIGUIENTE_MISMA_COHORTE']})
for c in closure:c['person']=pseudo.get(c.pop('key'),'SIN_LLAVE')
out={'raw_source':{'path':str(rawpath),'sha256':digest(rawpath),'rows':len(raw),'sheet':'BASE_RETENCION_MU'},'engine_source':{'path':str(engpath),'sha256':digest(engpath),'sheet':'02_COHORTE_ELEGIBILIDAD'},'sources':sources,'inventory':inventory,'origin':baseout,'closures':closure}
(TMP/'profile.json').write_text(json.dumps(out,ensure_ascii=False,default=str),encoding='utf-8')
print(json.dumps({'sources':sources,'base':{str(y):{'candidates':sum(b['year']==y for b in baseout),'states':dict(Counter(b['annual_state'] for b in baseout if b['year']==y)),'types_doc':dict(Counter(b['type_doc'] for b in baseout if b['year']==y)),'invalid_dv':sum(not b['dv_ok'] for b in baseout if b['year']==y)} for y in (2024,2025)}},ensure_ascii=False,indent=2))
