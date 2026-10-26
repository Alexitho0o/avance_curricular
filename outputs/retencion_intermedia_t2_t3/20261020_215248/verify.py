from pathlib import Path
import json, hashlib, zipfile, shutil, subprocess
from collections import defaultdict
import openpyxl

out=Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/retencion_intermedia_t2_t3/20261020_215248')
tmp=Path(__file__).resolve().parent
d=json.loads((tmp/'profile.json').read_text())
p=out/'REVISION_RETENCION_INTERMEDIA_IPSS_2024_2025_T2_T3_20261020.xlsx'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
sources=d['sources']+[d['raw_source'],d['engine_source']]
for src in sources:assert sha(src['path'])==src['sha256'],src['path']
assert d['raw_source']['sha256']=='19cf33ae4925a7c13ec4745394141a297f4f64be4052a240cb0d2c42ffeb53bd'
with zipfile.ZipFile(p) as z:assert z.testzip() is None
w=openpyxl.load_workbook(p,read_only=False,data_only=False)
errors=[];formulas=0
for s in w:
    for row in s:
        for c in row:
            formulas+=c.data_type=='f'
            if c.data_type=='e':errors.append((s.title,c.coordinate,c.value))
    for t in s.tables.values():
        a,b,c,e=openpyxl.utils.range_boundaries(t.ref)
        heads=[s.cell(b,k).value for k in range(a,c+1)]
        assert heads==[x.name for x in t.tableColumns],t.name
assert not errors,errors
assert len(w.sheetnames)==6
assert len(w['Resumen']._charts)==1
assert w['Resumen'].data_validations.dataValidation[0].formula1=='ListaCarreras'
assert 'ListaCarreras' in w.defined_names
w.close()
w=openpyxl.load_workbook(p,read_only=True,data_only=True)
idx=defaultdict(list)
for x in d['closures']:idx[(x['source'],x['person'])].append(x)
detail=list(w['Cruce_revision'].iter_rows(min_row=7,max_row=1897,values_only=True))
for b,row in zip(d['origin'],detail):
    for term,offset in [(2,9),(3,13)]:
        found=idx[(f'T{term}_{b["year"]+1}',b['person'])] if b['key_ok'] else []
        exp=[len(found),sum(x['state']=='VIGENTE' for x in found),sum(x['state'] not in ('VIGENTE','NO INFORMADO') for x in found),sum(x['state']=='NO INFORMADO' for x in found)]
        assert list(row[offset:offset+4])==exp,(b['source_row'],term)
for row in w['Por_carrera'].iter_rows(min_row=7,max_row=34,values_only=True):
    y,career=row[:2];b=[x for x in d['origin'] if x['year']==y and x['career']==career];inc=[x for x in b if x['annual_eligible']];key=[x for x in inc if x['key_ok']]
    exp=[len(b),len(inc),len(inc)-len(key)]
    for term in (2,3):
        matches=[x for x in key if idx[(f'T{term}_{y+1}',x['person'])]]
        exp.extend([len(matches),'n.d.' if y==2025 and term==2 else sum(any(z['state']=='VIGENTE' for z in idx[(f'T{term}_{y+1}',x['person'])]) for x in matches)])
    assert list(row[2:9])==exp,(y,career,row[2:9],exp)
for s in w:
    for row in s:
        for cell in row:
            if cell.data_type=='e':errors.append((s.title,cell.coordinate,cell.value))
assert not errors
w.close()
manifest=json.loads((out/'MANIFEST_REVISION.json').read_text())
assert manifest['salida']['sha256']==sha(p)
head=subprocess.check_output(['git','-C','/Users/alexi/Documents/GitHub/avance_curricular','rev-parse','HEAD'],text=True).strip()
branch=subprocess.check_output(['git','-C','/Users/alexi/Documents/GitHub/avance_curricular','branch','--show-current'],text=True).strip()
manifest['git'].update({'rama_despues':branch,'head_despues':head})
manifest['qa']={'resultado':'OK revisión; tasas pendientes','formula_errors':0,'formula_cells':formulas,'detail_rows_checked':len(detail),'career_rows_checked':28,'sheets':6,'charts':1,'sources_hashes_unchanged':True,'excel_native_open_test':'No realizado; recálculo y selector verificados con artifact-tool y lectura independiente del XLSX exportado.'}
for name in ('profile.py','build.mjs','verify.py'):
    target=out/name
    if (tmp/name).resolve()!=target.resolve():shutil.copy2(tmp/name,target)
manifest['reproducibilidad']={'pasos':['Ejecutar profile.py con Python y openpyxl de lectura: regenera profile.json seudonimizado a partir de fuentes originales.','Ejecutar build.mjs con runtime @oai/artifact-tool: genera libro y comprueba fórmulas y selector.','Ejecutar verify.py: compara cada cruce y carrera contra cálculo independiente y hashes originales.'],'scripts':[{'archivo':name,'sha256':sha(out/name)} for name in ('profile.py','build.mjs','verify.py')],'datos_intermedios':'profile.json se genera localmente; no se entrega ni se agrega a Git.'}
(out/'MANIFEST_REVISION.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
inspection=out/(p.name+'.inspect.ndjson')
if inspection.exists():inspection.unlink()
print(json.dumps({'file':str(p),'sha256':sha(p),'qa':manifest['qa'],'git':manifest['git']},ensure_ascii=False,indent=2))
