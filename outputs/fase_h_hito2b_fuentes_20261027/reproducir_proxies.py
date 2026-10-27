from pathlib import Path
from collections import defaultdict
import openpyxl,json,hashlib,os,sqlite3
O=Path(__file__).parent;P=Path.home()/'.local/share/ipss/fase_h_hito2b_20261027';v=P.parent/'fase_h_regla_intertrimestral_20261027/V24_QA_COPIA.xlsx';prom=Path('/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_12166.xlsx');sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();assert sha(v)=='1e28d5955d874e7d86ec0ca87ab6a2cc478515c0dbd604f4fc86f60fa491c481';assert sha(prom)=='aad1a71d166bdfb839049a5db534a21d050df1ed4e9251ad9982245896429015'
w=openpyxl.load_workbook(v,read_only=True,data_only=True);people=defaultdict(set)
for r in w['20_Foto_BBDD'].iter_rows(min_row=6,max_row=4478,values_only=True):
 if r[2]==1 and r[1] in (2024,2026) and r[8] is not None:people[r[1]].add(str(int(r[8])))
z=openpyxl.load_workbook(prom,read_only=True,data_only=True);it=z['Hoja1'].iter_rows(values_only=True);hd={str(v):i for i,v in enumerate(next(it)) if v is not None};found=defaultdict(set);observed=defaultdict(set);evidence=defaultdict(list)
for rn,row in enumerate(it,2):
 def val(k):return row[hd[k]]
 try:body=str(int(val('RUT')));year=int(val('ANO'));per=int(val('PERIODO'))
 except (ValueError,TypeError):continue
 if (year,per) not in ((2025,2),(2025,3),(2026,2)):continue
 coh=2024 if year==2025 else 2026
 if body not in people[coh]:continue
 observed[(coh,year,per)].add(body)
 if str(val('ESTADO_ACADEMICO')).strip().upper()=='VIGENTE':found[(coh,year,per)].add(body);evidence[(coh,body,year,per)].append(rn)
z.close();s=w['26_Avance_Curricular_Resumen'];records=[]
for r,coh,yr,per in [(1850,2024,2025,2),(1851,2024,2025,3),(1854,2026,2026,2)]:
 key=(coh,yr,per);records.append({'celda':'E'+str(r),'personas_en_cohorte':len(people[coh]),'actividad_observada':len(observed[key]),'vigente_extracto_reproducido':len(found[key]),'valor_literal_v24':s['E'+str(r)].value,'coincide':len(found[key])==s['E'+str(r)].value})
w.close();assert all(r['coincide'] for r in records);db=sqlite3.connect(P/'PROXIES_PRIVADOS.sqlite');db.execute('create table actividad (cohorte integer, clave_cuerpo text, anio integer, periodo integer, filas_promedios text)')
for (coh,body,year,per),rows in evidence.items():db.execute('insert into actividad values(?,?,?,?,?)',(coh,body,year,per,json.dumps(rows)))
db.commit();db.close();res={'fuente_v24':{'ruta':str(v),'sha256':sha(v)},'fuente_promedios':{'ruta':str(prom),'sha256':sha(prom),'hoja':'Hoja1'},'resultados':records,'metodo':'Reconstrucción por cuerpo de RUT, periodo y estado VIGENTE del extracto; no demuestra que el constructor original usara el mismo algoritmo.','limites':['DV no fue usado en el método heredado; no acredita identidad de forma universal.','Estado del extracto no demuestra estado a fecha de cierre.'],'detalle_privado':{'ruta':str(P/'PROXIES_PRIVADOS.sqlite'),'sha256':sha(P/'PROXIES_PRIVADOS.sqlite')}};(O/'PROXIES_REPRODUCIDOS.json').write_text(json.dumps(res,ensure_ascii=False,indent=2)+'\n');print(json.dumps(records,ensure_ascii=False))
