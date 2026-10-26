"""C: evidencia agregada, edición acotada y QA de tres tasas; fuentes solo lectura."""
import argparse, copy, hashlib, json, re, shutil, subprocess, zipfile
from datetime import datetime
from zoneinfo import ZoneInfo
from collections import Counter
from pathlib import Path
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
import openpyxl

R = Path(__file__).resolve().parents[1]
O = R/'outputs/aclaracion_tasas_2026_20261023_v16'
SRC = R/'outputs/correccion_proxy_t3_2025_20261023_v15/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V15.xlsx'
ORIGINAL = Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
NAME = 'RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V16.xlsx'
SHEET = '18_Foto_Intertrimestral'
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q = lambda s:'{'+NS+'}'+s
sha = lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(n, x):
    (O/n).write_text(json.dumps(x, ensure_ascii=False, indent=2)+'\n')
def git():
    return {k:subprocess.check_output(['git', *v],cwd=R,text=True).strip() for k,v in [('rama',['branch','--show-current']),('HEAD',['rev-parse','HEAD'])]}

def prepare():
    import unicodedata
    t2 = [p for p in Path('/Users/alexi/Downloads').glob('*.xlsx') if unicodedata.normalize('NFC',p.name)=='Matrícula T_2_13_07_2026.xlsx']
    assert len(t2)==1
    expected = {SRC:'3d5390c763e6ef40f327a87252a3793c185c4c4423c754b533fd96ec85abe3b3',ORIGINAL:'20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8',t2[0]:'7926a9cabcc67ed19abb50459132fd74ebb323bff9f5cb53e9116378c21a6769'}
    for p,h in expected.items(): assert sha(p)==h,str(p)
    O.mkdir(parents=True,exist_ok=True)
    w=openpyxl.load_workbook(ORIGINAL,read_only=True,data_only=True)
    rows=[r for r in w['20_Foto_BBDD'].iter_rows(min_row=6,max_row=4478,values_only=True) if r[1]==2026]
    groups=[r for r in rows if r[2]==1]
    unique=[r for r in rows if r[19]==1]
    stats=lambda rr:dict(personas=len(rr),vigentes=sum(r[16] for r in rr),base=sum(r[17] for r in rr),excluidos=sum(r[17]==0 for r in rr),estados=dict(Counter(r[15] for r in rr)))
    a,b=stats(groups),stats(unique)
    assert [a[k] for k in ['personas','vigentes','base','excluidos']]==[1471,1360,1467,4]
    assert [b[k] for k in ['personas','vigentes','base','excluidos']]==[2523,2290,2519,4]
    keys={r[14] for r in unique if r[17]==0}
    cuts=[r for r in w['25_Foto_Corte_T3_2026'].iter_rows(min_row=6,values_only=True) if len(r)>12 and r[12] in keys]
    states=dict(Counter(str(r[9])+' | '+str(r[10]) for r in cuts))
    assert states=={'TITULADO | 31 - TITULADO APROBADO':4},states
    wt=openpyxl.load_workbook(t2[0],read_only=True,data_only=True)
    data=list(wt['Hoja1'].iter_rows(values_only=True))
    headers=[str(x).strip().upper() if x is not None else '' for x in data[0]]
    assert 'CODCLI' in headers and 'RUT' in headers
    rr=[r for r in data[1:] if r[headers.index('CODCLI')] is not None]
    totals=dict(matriculas=len(rr),codcli=len({r[headers.index('CODCLI')] for r in rr}),personas=len({r[headers.index('RUT')] for r in rr}))
    assert totals==dict(matriculas=528,codcli=528,personas=527),totals
    v=openpyxl.load_workbook(SRC,read_only=True,data_only=True)
    rs=list(v['23_Foto_Renuncias_T2_2026'].iter_rows(min_row=6,max_row=24,values_only=True))
    t=[sum(r[i] for r in rs) for i in [2,3,5]]
    assert t==[528,66,462],t
    s53='53_Codigos_T1_Confirmados'
    rr53=list(v[s53].iter_rows(values_only=True)); end=len(rr53)
    filtered=[r for r in rr53[5:] if r[:3]==(2026,2026,3)]
    assert [sum(r[i] for r in filtered) for i in [5,6,7]]==[1471,1360,1467]
    wf=openpyxl.load_workbook(ORIGINAL,read_only=True,data_only=False)
    oldformulas={c:wf[SHEET][c].value for c in ['H11','F12','H14']}
    evidence=dict(clasificacion='B/C; exclusión según decisión interna D existente, no norma SIES',fuentes={str(p):{'sha256':h,'tipo':'fuente institucional' if p==t2[0] else 'resultado/derivado'} for p,h in expected.items()},T1_T3=a,total_unico_T3=b,exclusiones=states,T2=totals|dict(renuncias=66,vigentes=462),formulas_originales=oldformulas,referencias=['20_Foto_BBDD!B6:V4478: B año, C grupo, P estado, Q numerador, R elegibilidad, T deduplicación','25_Foto_Corte_T3_2026!J6:M15190: estado/situación y llave de matrícula','21_Foto_Reglas_VIG!B3 y B31:E31: regla interna TITULADO','23_Foto_Renuncias_T2_2026!C6:F24: matrícula, renuncias, diferencia','T2 original: Hoja1 filas 2:529; Hoja2 B32:C50',f'{s53}!A6:H{end}: cohorte, año, período, población, numerador, denominador'],pendiente_e='Renuncias individuales para calcular una tasa T2 de personas. La tasa vigente usa matrículas y se conserva.')
    dump('EVIDENCIA_AGREGADA.json',evidence)
    edits=[]
    def add(cell,value=None,formula=None,expected=None,style_from=None):
        e=dict(sheet=SHEET,cell=cell,before=v[SHEET][cell].value)
        e.update(dict(value=value,expected=value) if formula is None else dict(formula=formula,expected=expected))
        if style_from:e['style_from']=style_from
        edits.append(e)
    add('B91','Bases de cálculo de las tasas 2026',style_from='B79')
    for c,label in [('B','Medida'),('D','Personas del grupo'),('E','Titulados excluidos'),('F','Base de la tasa'),('G','Vigentes'),('H','Tasa'),('I','Unidad de la tasa')]:
        add(c+'92',label,style_from=c+'6')
    for row,label in [(93,'T1 · cierre T3'),(94,'Ingreso T2 · cierre T2'),(95,'Total único · cierre T3')]:
        add('B'+str(row),label,style_from='C11')
        add('I'+str(row),'Matrículas' if row==94 else 'Personas',style_from='C11')
    fs={
        'D93':('=D11',1471),'F93':(f'=SUMIFS(\'{s53}\'!$H$6:$H${end},\'{s53}\'!$A$6:$A${end},2026,\'{s53}\'!$B$6:$B${end},2026,\'{s53}\'!$C$6:$C${end},3)',1467),'E93':('=D93-F93',4),'G93':('=G11',1360),
        'D94':('=D12',527),'F94':("=SUM('23_Foto_Renuncias_T2_2026'!$C$6:$C$24)",528),'G94':('=E12',462),
        'D95':('=D14',2523),'F95':('=D95-E95',2519),'G95':('=G14',2290)}
    for c,(f,n) in fs.items():add(c,formula=f,expected=n,style_from='D11')
    add('E94','No aplica',style_from='C11');add('E95',4,style_from='D11')
    for row,n,d in [(93,1360,1467),(94,462,528),(95,2290,2519)]:
        add('H'+str(row),formula=f'=IF(F{row}>0,G{row}/F{row},"—")',expected=n/d,style_from='H11')
    add('H11',formula='=H93',expected=1360/1467)
    add('F12',formula='=H94',expected=462/528)
    add('H14',formula='=H95',expected=2290/2519)
    add('B96','B/C/D · T3 excluye a cuatro TITULADO / 31 - TITULADO APROBADO. Se conserva la regla interna de la hoja 21. Fuente: original auditado 22-10-2026, hoja 20 (P/Q/R/T); T1 conciliado con hoja 53.',style_from='B80')
    add('B97','B/C · T2: 528 matrículas de 527 personas; 66 renuncias agregadas. 462/528 mide matrículas (hoja 23). E · Para una tasa de personas faltan las renuncias individuales; no se divide 462/527.',style_from='B80')
    dump('EDICION.json',dict(source=str(SRC),source_sha256=sha(SRC),scratch='/tmp/retencion_tasas_2026_v16',edits=edits,merges=['B91:I91','B92:C92','B93:C93','B94:C94','B95:C95','B96:I96','B97:I97'],heights={91:21,92:42,93:30,94:30,95:30,96:36,97:36},git_antes=git(),fuentes=evidence['fuentes']))
    print('Evidencia reproducida y',len(edits),'celdas preparadas; sin identificadores.')

def finish():
    m=json.loads((O/'EDICION.json').read_text());assert sha(SRC)==m['source_sha256']
    z=zipfile.ZipFile(SRC);a=zipfile.ZipFile(Path(m['scratch'])/'artifact_recalculado.xlsx')
    names=[s.get('name') for s in ET.fromstring(z.read('xl/workbook.xml')).find(Q('sheets'))]
    part=f'xl/worksheets/sheet{names.index(SHEET)+1}.xml';txt=z.read(part).decode()
    before=ET.fromstring(txt);old={c.get('r'):c for c in before.iter(Q('c'))}
    acells={c.get('r'):c for c in ET.fromstring(a.read(part)).iter(Q('c'))}
    strings=[''.join(t.text or '' for t in n.iter(Q('t'))) for n in ET.fromstring(a.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in a.namelist() else []
    newrows={}
    for e in m['edits']:
        address=e['cell'];c=acells[address];attrs={'r':address}
        ref=old.get(address,old.get(e.get('style_from')))
        if ref is not None:attrs.update(ref.attrib);attrs['r']=address
        attrs.pop('t',None)
        if c.get('t'):attrs['t']=c.get('t')
        if c.get('t')=='s':
            attrs['t']='inlineStr';body='<is><t xml:space="preserve">'+escape(strings[int(c.find(Q('v')).text)])+'</t></is>'
        else:
            body=''.join(ET.tostring(n,encoding='unicode').replace('ns0:','').replace(':ns0','') for n in c)
            body=re.sub(r' xmlns="[^"]+"','',body)
        new='<c '+' '.join(k+'="'+escape(v,{'"':'&quot;'})+'"' for k,v in attrs.items())+'>'+body+'</c>'
        hit=re.search(rf'<c\b(?=[^>]*\br="{address}")[^>]*>.*?</c>',txt,re.S)
        if hit:txt=txt[:hit.start()]+new+txt[hit.end():]
        else:newrows.setdefault(int(re.search(r'\d+',address).group()),[]).append(new)
    col_order=lambda cell:openpyxl.utils.column_index_from_string(re.search(r' r="([A-Z]+)\d+"',cell).group(1))
    rows=''.join(f'<row r="{r}" ht="{m["heights"][str(r)]}" customHeight="1">'+''.join(sorted(cs,key=col_order))+'</row>' for r,cs in sorted(newrows.items()))
    txt=txt.replace('</sheetData>',rows+'</sheetData>')
    txt=re.sub(r'<dimension ref="[^"]+"\s*/>','<dimension ref="A1:I97"/>',txt)
    merges=before.find(Q('mergeCells'));assert merges is not None
    n=len(merges)+len(m['merges'])
    txt=re.sub(r'<mergeCells[^>]*>',f'<mergeCells count="{n}">',txt)
    txt=txt.replace('</mergeCells>',''.join('<mergeCell ref="'+r+'"/>' for r in m['merges'])+'</mergeCells>')
    dst=O/NAME
    # Idempotent rebuild of this derivative only; sources remain immutable.
    with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
        for info in z.infolist():out.writestr(copy.copy(info),txt.encode() if info.filename==part else z.read(info.filename))
    zz=zipfile.ZipFile(dst);assert zz.testzip() is None
    for p in zz.namelist():
        if p.endswith(('.xml','.rels')):ET.fromstring(zz.read(p))
    assert z.namelist()==zz.namelist()
    assert [p for p in z.namelist() if z.read(p)!=zz.read(p)]==[part]
    wf=openpyxl.load_workbook(dst,read_only=True,data_only=False);wv=openpyxl.load_workbook(dst,read_only=True,data_only=True)
    pv=openpyxl.load_workbook(SRC,read_only=True,data_only=True)
    checks=[]
    for e in m['edits']:
        assert wf[SHEET][e['cell']].value==e.get('formula',e.get('value')),e['cell']
        v=wv[SHEET][e['cell']].value;want=e['expected']
        assert abs(v-want)<1e-12 if isinstance(want,(float,int)) else v==want,(e['cell'],v,want)
        checks.append(dict(celda=e['cell'],esperado=want,observado=v))
    for rr in pv[SHEET].iter_rows(max_row=89):
        for c in rr:
            if not hasattr(c,'coordinate'):continue
            v=wv[SHEET][c.coordinate].value
            assert abs(v-c.value)<1e-12 if isinstance(c.value,(float,int)) and not isinstance(c.value,bool) else v==c.value,c.coordinate
    # Every original cell XML remains unchanged, except three authorized formulas.
    final={c.get('r'):c for c in ET.fromstring(zz.read(part)).iter(Q('c'))}
    for row in ET.fromstring(zz.read(part)).find(Q('sheetData')):
        addresses=[openpyxl.utils.cell.coordinate_from_string(c.get('r'))[0] for c in row if c.tag==Q('c')]
        indices=[openpyxl.utils.column_index_from_string(c) for c in addresses]
        assert indices==sorted(set(indices)),('orden o duplicación de celdas',row.get('r'))
    changed=[]
    for c,oldc in old.items():
        if ET.tostring(oldc)!=ET.tostring(final[c]):changed.append(c)
    assert set(changed)=={'H11','F12','H14'},changed
    charts=[p for p in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',p)]
    tables=[p for p in z.namelist() if re.fullmatch(r'xl/tables/table\d+\.xml',p)]
    assert (len(names),len(charts),len(tables))==(54,13,35)
    for p,h in m['fuentes'].items():assert sha(p)==h['sha256']
    assert git()==m['git_antes']
    dump('QA_ESTRUCTURA_Y_CALCULO.json',dict(estado='OK_PENDIENTE_EXCEL_NATIVO',controles=checks,partes_modificadas=[part],celdas_preexistentes_modificadas=changed,todos_los_valores_previos_conservados=True,hojas_conservadas_bytes=53,graficos_conservados_bytes=13,tablas_conservadas_bytes=35,estilos_conservados_bytes=True,fuentes_y_denominadores_conservados=True,git_antes=m['git_antes'],git_despues=git(),sha256=sha(dst)))
    print(dst)

def close():
    m=json.loads((O/'EDICION.json').read_text())
    qa=json.loads((O/'QA_ESTRUCTURA_Y_CALCULO.json').read_text())
    dyn=json.loads((O/'QA_DINAMICA.json').read_text())
    native=json.loads((O/'QA_EXCEL_NATIVO.json').read_text())
    assert dyn['estado']=='OK' and native['apertura_sin_reparacion'] is True
    dst=O/NAME;assert sha(dst)==qa['sha256']
    for p,h in m['fuentes'].items():assert sha(p)==h['sha256']
    assert git()==m['git_antes']
    desktop=Path('/Users/alexi/Desktop')/NAME
    if desktop.exists():assert sha(desktop)==sha(dst),'No sobrescribir una copia distinta del usuario'
    else:shutil.copy2(dst,desktop)
    assert sha(dst)==sha(desktop)
    w=openpyxl.load_workbook(dst,read_only=True,data_only=True)
    errors=[c.coordinate for rr in w[SHEET].iter_rows(min_row=91,max_row=97) for c in rr if c.data_type=='e']
    assert not errors
    status=subprocess.check_output(['git','status','--short'],cwd=R,text=True)
    initial=Path('/tmp/git_antes_tasas_v16.txt')
    before=initial.read_text() if initial.exists() else ''
    delta=sorted(set(status.splitlines())-set(before.splitlines()))
    dump('GIT_ESTADO_ENTREGA.json',dict(antes=m['git_antes'],despues=git(),commit=False,push=False,nuevas_lineas_desde_inicio=delta,limite='El repositorio ya tenía cambios ajenos; no se incorporaron ni alteraron.'))
    now=datetime.now(ZoneInfo('America/Santiago')).isoformat()
    qa['estado']='OK_ALCANCE_TRES_TASAS';qa['errores_nuevas_formulas']=errors
    dump('QA_ESTRUCTURA_Y_CALCULO.json',qa)
    readme='''# V16 · Tasas 2026 de la hoja 18

B/C: se reprodujeron las tres tasas a partir del original auditado y las bases de cálculo vigentes. No se encontró una división incorrecta. V15 había conservado estos resultados como constantes tras la agregación; V16 devuelve fórmulas trazables y hace explícitas las unidades y exclusiones en B91:I97.

| Celda | Población visible | Base efectiva | Numerador | Tasa | Respaldo |
|---|---:|---:|---:|---:|---|
| H11 | 1.471 personas | 1.467 personas elegibles | 1.360 | 92,706203 % | B/C; cuatro titulados excluidos según D existente |
| F12 | 527 personas | 528 matrículas | 462 | 87,5 % | B/C; 528 matrículas menos 66 renuncias |
| H14 | 2.523 personas únicas | 2.519 personas elegibles | 2.290 | 90,909091 % | B/C; mismas cuatro exclusiones D, deduplicación anual |

D: la exclusión NO APLICA de TITULADO / 31 - TITULADO APROBADO ya está declarada en la hoja 21 y las notas originales. Se preservó; este trabajo no la presenta como norma SIES.

E: para transformar el indicador T2 de matrículas en uno de personas hacen falta renuncias individuales. No se infiere esa tasa ni se cambia 528 por 527. Las tres tasas revisadas conservan sus valores. Los otros pendientes de V15 permanecen fuera de este alcance.

## Reproducción y controles

C: ejecutar el script Python `aclarar_tasas_2026_v16_20261023.py prepare`; autoría con `editar_tasas_2026_v16_20261023.mjs <salida> edit`; ensamblar con Python `finish`; verificar dinámica con Node `qa`; abrir el resultado en Excel y registrar QA_EXCEL_NATIVO.json antes de Python `close`. Se usan las dependencias del runtime de Codex y las rutas/hashes de EDICION.json. La fuente individual se lee solo en memoria y solo se persisten agregados.

C: los cambios se limitan a tres fórmulas preexistentes y al nuevo bloque B91:I97. Las otras 53 hojas, 13 gráficos, 35 tablas y estilos mantienen sus bytes originales. Los valores previos y denominadores no cambiaron. Se probaron las dependencias con cambios temporales en memoria, incluido denominador cero. El archivo final abrió en Excel sin aviso de reparación. El primer intento de ensamblaje tenía celdas nuevas fuera de orden; se corrigió y se agregó control estricto de orden y duplicados antes de la entrega.

C: el QA se limita a estas tasas y a la preservación del resto; no equivale a una nueva auditoría funcional de las 54 hojas ni al cierre de pendientes históricos. No se guardó el libro desde Excel: se conservó la estructura del derivado verificado.

## Checkpoint

C: V16 lista para revisión local. Bases, numeradores, denominadores, porcentajes y diccionario conservados. Copia del Escritorio idéntica. Sin commit, push ni publicación de datos personales. Evidencia agregada y hashes en MANIFEST_TASAS_2026_V16.json. Si se cambia a futuro la unidad o la regla de exclusión, será otra decisión explícita y otra versión; no debe aplicarse silenciosamente.
'''
    (O/'README_Y_REANUDACION.md').write_text(readme)
    files=[p for p in O.iterdir() if p.is_file() and p.name!='MANIFEST_TASAS_2026_V16.json']
    support=[R/'scripts/aclarar_tasas_2026_v16_20261023.py',R/'scripts/editar_tasas_2026_v16_20261023.mjs',R/'docs/ESPECIFICACION_BASES_TASAS_2026_V16_20261023.md']
    manifest=dict(hito='H001-RETENCION-IPSS162',proceso='seguimiento interno intertrimestral',version='V16',fecha_ejecucion=now,periodo_revisado=2026,alcance='18_Foto_Intertrimestral H11/F12/H14',estado='QA_OK_LOCAL_SIN_COMMIT_PUSH',clasificaciones={'A':'No se incorpora ni modifica regla oficial','B':'Datos fuente y exclusiones observados; evidencia agregada','C':'Fórmulas, conciliación, preservación OOXML y QA','D':'Se conserva exclusión interna ya documentada; sin nueva regla','E':'Tasa T2 por persona no calculable con renuncias solo agregadas'},fuentes=m['fuentes'],resultados=[dict(celda=c,numerador=n,denominador=d,tasa=n/d) for c,n,d in [('H11',1360,1467),('F12',462,528),('H14',2290,2519)]],sin_cambio_bases_ni_denominadores=True,archivos={str(p):{'sha256':sha(p),'bytes':p.stat().st_size} for p in files+support},copia_escritorio={'ruta':str(desktop),'sha256':sha(desktop)},qa=dict(estructura='QA_ESTRUCTURA_Y_CALCULO.json',dinamica='QA_DINAMICA.json',nativo='QA_EXCEL_NATIVO.json',limite='Tres tasas; preservación byte a byte del contenido ajeno a la edición, no nueva certificación global'),privacidad='Solo agregados nuevos, sin identificadores individuales',git_antes=m['git_antes'],git_despues=git(),commit=False,push=False)
    dump('MANIFEST_TASAS_2026_V16.json',manifest)
    print(json.dumps(dict(salida=str(dst),escritorio=str(desktop),sha256=sha(dst),rama=git()['rama'],HEAD=git()['HEAD'],cambios_nuevos=delta),ensure_ascii=False))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('fase',choices=['prepare','finish','close']);args=p.parse_args()
    {'prepare':prepare,'finish':finish,'close':close}[args.fase]()
