"""Contrasta INF35-03 en respaldo anual 2024 enlazado por matrícula incremental."""
from pathlib import Path
import json,sys
from openpyxl import load_workbook
sys.path.insert(0,str(Path.cwd()))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado
O=Path('outputs/origen_v1_inf35_03_20261023_150536');cfg=json.loads((O/'CONFIG.json').read_text());P=Path(cfg['privado']);e=json.loads((P/'EVIDENCIA.json').read_text());person=e['objetivos'][0]['persona']
inv=json.loads(Path('outputs/retencion_intermedia_t2_t3/20261021_093412_git_onedrive_mu2025/INVENTARIO_ONEDRIVE.json').read_text());p=next(Path(x['ruta']) for x in inv['archivos'] if Path(x['ruta']).name=='Matrícula_Pregrado_2024.xlsx');h=sha256(p);w=load_workbook(p,read_only=True,data_only=True);out=[]
for sh,start in [('Intake 1',3),('08-05-2024_18-31-54_Matrícula_P',1),('Todos',2)]:
 s=w[sh];hdr=next(s.iter_rows(min_row=start,max_row=start,max_col=38,values_only=True));i=hdr.index('N_DOC');dv=hdr.index('DV')
 for rn,row in enumerate(s.iter_rows(min_row=start+1,max_col=38,values_only=True),start+1):
  try:key=rut_normalizado(str(int(row[i]))+'-'+str(row[dv]))
  except (ValueError,TypeError):continue
  if key!=person:continue
  out.append({'hoja':sh,'fila':rn,'valores':{k:row[j] for j,k in enumerate(hdr) if k in ['COD_IES','COD_SED','COD_CAR','JOR','VERSION','ANIO_ING_ACT','ANIO_ING_ORI','FECH_CAR','FECHA_MATRICULA']}})
w.close();assert sha256(p)==h
z={'fuente':str(p),'sha256':h,'coincidencias':out,'limite':'Retiene información de la matrícula anual; nombre del respaldo no demuestra jornada T3.'};(O/'RESPALDO_2024.json').write_text(json.dumps(z,ensure_ascii=False,indent=2,default=str)+'\n');print(json.dumps(z,ensure_ascii=False,default=str))
