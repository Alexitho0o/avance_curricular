#!/usr/bin/env python3
"""Construye el reporte gerencial definitivo de brechas MU SIES 2026.

El script no modifica fuentes originales. Genera un workbook estático de
cuatro hojas, archivos de validación, manifiesto e inventario de hashes.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
SOURCE_BOOK = Path("/Users/alexi/Desktop/AUDITORIA_MATRICULA_UNIFICADA_2026_RECONSTRUIDA_PARA_REVISION_20260731_002214.xlsx")
CSV_AUDIT_DIR = REPO / "auditoria_comparacion_csv_mu2026" / "2026-07-31_114529"
DOWNLOADS = Path("/Users/alexi/Downloads")
OUT_ROOT = REPO / "outputs" / "reporte_gerencial_brechas_mu2026"
DESKTOP_FINAL = Path.home() / "Desktop" / "REPORTE_GERENCIAL_BRECHAS_MATRICULA_UNIFICADA_2026_DEFINITIVO.xlsx"

PREGRADO_SHEET = "08_PREGRADO_COMPLETO"
PUBLICACION_SHEET = "07_PUBLICACION_COMPLETA"
CRUCE_SHEET = "02_CRUCE_66_CODIGOS"
POSGRADO_SHEET = "09_POSGRADO_COMPLETO"

OFFICIAL_PREGRADO_FIELDS = [
    "TIPO_DOC",
    "N_DOC",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRE",
    "SEXO",
    "FECH_NAC",
    "NAC",
    "PAIS_EST_SEC",
    "COD_SED",
    "COD_CAR",
    "MODALIDAD",
    "JOR",
    "VERSION",
    "FOR_ING_ACT",
    "ANIO_ING_ACT",
    "SEM_ING_ACT",
    "ANIO_ING_ORI",
    "SEM_ING_ORI",
    "ASI_INS_ANT",
    "ASI_APR_ANT",
    "PROM_PRI_SEM",
    "PROM_SEG_SEM",
    "ASI_INS_HIS",
    "ASI_APR_HIS",
    "NIV_ACA",
    "SIT_FON_SOL",
    "SUS_PRE",
    "FECHA_MATRICULA",
    "REINCORPORACION",
    "VIG",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def file_info(path: Path) -> dict[str, Any]:
    return {
        "ruta": str(path),
        "nombre": path.name,
        "tamano_bytes": path.stat().st_size if path.exists() else None,
        "fecha_modificacion": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds") if path.exists() else None,
        "sha256": sha256(path) if path.exists() else None,
    }


def rows_from_ws(ws) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def table_from_sheet(ws) -> tuple[list[str], list[dict[str, Any]]]:
    rows = rows_from_ws(ws)
    headers = [str(v) if v is not None else "" for v in rows[0]]
    data = []
    for raw in rows[1:]:
        data.append({headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))})
    return headers, data


def norm_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def as_int(v: Any) -> int:
    if v is None or v == "":
        return 0
    return int(float(v))


def codigo_from_row(row: dict[str, Any]) -> str:
    existing = norm_text(row.get("AUD_CODIGO_NORMALIZADO") or row.get("AUD_TV"))
    if existing and not existing.startswith("[FÓRMULA"):
        return existing
    return f"I162S{norm_text(row.get('COD_SED'))}C{norm_text(row.get('COD_CAR'))}J{norm_text(row.get('JOR'))}V{norm_text(row.get('VERSION'))}"


def classify_breach(publicacion: int, institucional: int, estado: str) -> str:
    if estado == "SOLO_PUBLICACION":
        return "SOLO_PUBLICACION"
    if publicacion > institucional:
        return "PUBLICACION_MAYOR"
    if publicacion < institucional:
        return "INSTITUCIONAL_MAYOR"
    return "IGUAL"


def get_git(cmd: list[str]) -> str:
    try:
        return subprocess.check_output(["git", "-C", str(REPO), *cmd], text=True).strip()
    except Exception:
        return "NO_DISPONIBLE"


def style_title(ws, cell: str, value: str):
    ws[cell] = value
    ws[cell].font = Font(bold=True, size=18, color="17365D")
    ws[cell].alignment = Alignment(vertical="center")


def style_section(ws, row: int, title: str, start_col: int = 1, end_col: int = 8):
    ws.cell(row=row, column=start_col, value=title)
    ws.cell(row=row, column=start_col).font = Font(bold=True, color="FFFFFF", size=12)
    ws.cell(row=row, column=start_col).fill = PatternFill("solid", fgColor="17365D")
    ws.cell(row=row, column=start_col).alignment = Alignment(vertical="center")
    if end_col > start_col:
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
    ws.row_dimensions[row].height = 22


def write_table(ws, start_row: int, start_col: int, headers: list[str], rows: list[list[Any]], table_name: str,
                style_name: str = "TableStyleMedium2") -> tuple[int, str]:
    for j, h in enumerate(headers, start_col):
        cell = ws.cell(start_row, j, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, start_row + 1):
        for j, value in enumerate(row, start_col):
            cell = ws.cell(i, j, value)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    end_row = start_row + len(rows)
    end_col = start_col + len(headers) - 1
    ref = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{max(end_row, start_row)}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=style_name, showFirstColumn=False, showLastColumn=False,
                                        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)
    ws.auto_filter.ref = ref
    return end_row, ref


def auto_width(ws, min_width: int = 9, max_width: int = 42):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 120) + 1):
            v = ws.cell(row, col).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[letter].width = max(min_width, min(max_len + 2, max_width))


def add_card(ws, row: int, col: int, title: str, value: Any, fill: str = "D9EAF7"):
    ws.cell(row, col, title)
    ws.cell(row, col).font = Font(bold=True, color="17365D")
    ws.cell(row, col).fill = PatternFill("solid", fgColor=fill)
    ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row + 1, col, value)
    ws.cell(row + 1, col).font = Font(bold=True, size=16, color="17365D")
    ws.cell(row + 1, col).fill = PatternFill("solid", fgColor="FFFFFF")
    ws.cell(row + 1, col).alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    for r in (row, row + 1):
        ws.cell(r, col).border = Border(left=thin, right=thin, top=thin, bottom=thin)


def signed(n: int) -> str:
    return f"{n:+,}".replace(",", ".")


def parse_xml_and_rels(path: Path) -> tuple[int, int, int]:
    xml_errors = 0
    external_rels = 0
    external_hyperlinks = 0
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            xml_errors += 1
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                try:
                    content = z.read(name)
                    root = ET.fromstring(content)
                    if name.endswith(".rels"):
                        for rel in root:
                            target_mode = rel.attrib.get("TargetMode", "")
                            target = rel.attrib.get("Target", "")
                            rtype = rel.attrib.get("Type", "")
                            if target_mode.lower() == "external" or re.match(r"^[a-z]+://", target, re.I):
                                external_rels += 1
                                if "hyperlink" in rtype.lower():
                                    external_hyperlinks += 1
                except Exception:
                    xml_errors += 1
    return xml_errors, external_rels, external_hyperlinks


def validate_formula_errors(path: Path) -> int:
    wb = load_workbook(path, read_only=True, data_only=False)
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!"}
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and any(err in v for err in errors):
                    count += 1
    return count


def main() -> int:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = OUT_ROOT / ts
    out_dir.mkdir(parents=True, exist_ok=False)
    xlsx_path = out_dir / f"REPORTE_GERENCIAL_BRECHAS_MATRICULA_UNIFICADA_2026_DEFINITIVO_{ts}.xlsx"
    validation_path = out_dir / "REPORTE_VALIDACION_REPORTE_GERENCIAL_MU2026.txt"
    manifest_path = out_dir / "MANIFEST_REPORTE_GERENCIAL_MU2026.json"
    inventory_path = out_dir / "INVENTARIO_FUENTES_SHA256.csv"

    expected_sources = [SOURCE_BOOK, CSV_AUDIT_DIR / "01_INVENTARIO_ARCHIVOS.csv", CSV_AUDIT_DIR / "03_COMPARACION_UNIVERSO.csv"]
    source_hashes_before = {str(p): sha256(p) for p in expected_sources if p.exists()}

    source_wb = load_workbook(SOURCE_BOOK, read_only=True, data_only=True)
    for sheet in [PUBLICACION_SHEET, PREGRADO_SHEET, POSGRADO_SHEET, CRUCE_SHEET]:
        if sheet not in source_wb.sheetnames:
            raise RuntimeError(f"Hoja requerida no existe: {sheet}")

    pub_headers, pub_rows = table_from_sheet(source_wb[PUBLICACION_SHEET])
    pre_headers_full, pre_rows_full = table_from_sheet(source_wb[PREGRADO_SHEET])
    pos_headers, pos_rows = table_from_sheet(source_wb[POSGRADO_SHEET])
    cruce_headers, cruce_rows = table_from_sheet(source_wb[CRUCE_SHEET])
    pre_headers = pre_headers_full[:32]
    pre_rows = [{h: row.get(h) for h in pre_headers} for row in pre_rows_full]

    if pre_headers != OFFICIAL_PREGRADO_FIELDS:
        raise RuntimeError("Las primeras 32 columnas de pregrado no coinciden con la estructura oficial esperada.")

    # Mapas de publicación: código, nombre y nivel observado.
    pub_by_code: dict[str, dict[str, Any]] = {}
    for r in pub_rows:
        code = norm_text(r.get("CÓDIGO CARRERA") or r.get("AUD_CODIGO_NORMALIZADO"))
        if not code:
            continue
        pub_by_code[code] = {
            "nombre": norm_text(r.get("NOMBRE CARRERA")) or "NO DISPONIBLE EN FUENTE",
            "nivel_global": norm_text(r.get("NIVEL GLOBAL")),
            "cantidad": as_int(r.get("TOTAL MATRÍCULA") or r.get("AUD_CANTIDAD")),
        }

    # Composición institucional de pregrado por código.
    inst_by_code: dict[str, dict[str, Any]] = defaultdict(lambda: {
        "vigente": 0,
        "vig0": 0,
        "vig1": 0,
        "vig2": 0,
        "rut_vigentes": set(),
        "sexo_vigente": Counter(),
        "cod_sed": "",
        "cod_car": "",
        "jor": "",
        "modalidad": "",
        "version": "",
    })
    total_vig = Counter()
    for r in pre_rows:
        code = codigo_from_row(r)
        d = inst_by_code[code]
        d["cod_sed"] = d["cod_sed"] or norm_text(r.get("COD_SED"))
        d["cod_car"] = d["cod_car"] or norm_text(r.get("COD_CAR"))
        d["jor"] = d["jor"] or norm_text(r.get("JOR"))
        d["modalidad"] = d["modalidad"] or norm_text(r.get("MODALIDAD"))
        d["version"] = d["version"] or norm_text(r.get("VERSION"))
        vig = norm_text(r.get("VIG"))
        total_vig[vig] += 1
        if vig == "0":
            d["vig0"] += 1
        elif vig == "1":
            d["vig1"] += 1
            d["vigente"] += 1
        elif vig == "2":
            d["vig2"] += 1
            d["vigente"] += 1
        if vig in {"1", "2"}:
            llave = "|".join([norm_text(r.get("TIPO_DOC")), norm_text(r.get("N_DOC")), norm_text(r.get("DV"))])
            d["rut_vigentes"].add(llave)
            sexo = norm_text(r.get("SEXO"))
            d["sexo_vigente"][sexo] += 1

    cruce_clean: list[dict[str, Any]] = []
    for r in cruce_rows:
        code = norm_text(r.get("CODIGO"))
        if not code:
            continue
        publicacion = as_int(r.get("PUBLICACION"))
        institucional = as_int(r.get("INSTITUCIONAL"))
        diferencia = publicacion - institucional
        estado = norm_text(r.get("ESTADO"))
        nivel = norm_text(r.get("NIVEL"))
        pinfo = pub_by_code.get(code, {})
        inst = inst_by_code.get(code, {})
        sexo_counter: Counter = inst.get("sexo_vigente", Counter()) if inst else Counter()
        hombres = sexo_counter.get("H", 0)
        mujeres = sexo_counter.get("M", 0)
        sin_info = sexo_counter.get("", 0)
        otras = {k: v for k, v in sexo_counter.items() if k not in {"H", "M", ""}}
        cruce_clean.append({
            "NIVEL": nivel,
            "CODIGO_SIES": code,
            "NOMBRE_CARRERA": pinfo.get("nombre") or "NO DISPONIBLE EN FUENTE",
            "PUBLICACION": publicacion,
            "INSTITUCIONAL_VIGENTE": institucional,
            "DIFERENCIA": diferencia,
            "ESTADO": estado,
            "TIPO_BRECHA": classify_breach(publicacion, institucional, estado),
            "RUT_UNICOS_INSTITUCIONALES_VIGENTES": len(inst.get("rut_vigentes", set())) if inst else 0,
            "HOMBRES_INSTITUCIONALES_VIGENTES": hombres,
            "MUJERES_INSTITUCIONALES_VIGENTES": mujeres,
            "OTRAS_CATEGORIAS_SEXO": "; ".join(f"{k}:{v}" for k, v in sorted(otras.items())) if otras else "",
            "SIN_INFORMACION_SEXO": sin_info,
            "REGISTROS_VIG_0": inst.get("vig0", 0) if inst else 0,
            "FUENTE_NOMBRE_CARRERA": "07_PUBLICACION_COMPLETA" if pinfo.get("nombre") else "NO DISPONIBLE EN FUENTE",
            "OBSERVACION": "",
        })

    diff_rows = [r for r in cruce_clean if r["DIFERENCIA"] != 0]
    diff_rows.sort(key=lambda r: (0 if r["DIFERENCIA"] > 0 else 1, -abs(r["DIFERENCIA"]), r["CODIGO_SIES"]))
    for i, r in enumerate(diff_rows, 1):
        r["ORDEN_BRECHA"] = i
        if r["TIPO_BRECHA"] == "SOLO_PUBLICACION":
            r["OBSERVACION"] = "Código presente solo en publicación; no existe registro vigente institucional asociado en el consolidado."
        elif r["DIFERENCIA"] > 0:
            r["OBSERVACION"] = "Publicación mayor que registros institucionales vigentes comparables."
        else:
            r["OBSERVACION"] = "Institucional vigente mayor que publicación."

    positives = [r for r in diff_rows if r["DIFERENCIA"] > 0]
    negatives = [r for r in diff_rows if r["DIFERENCIA"] < 0]
    iguales = [r for r in cruce_clean if r["DIFERENCIA"] == 0]
    solo_publicacion = [r for r in cruce_clean if r["ESTADO"] == "SOLO_PUBLICACION"]
    solo_institucional = [r for r in cruce_clean if r["ESTADO"] == "SOLO_INSTITUCIONAL"]

    pub_pre = sum(r["PUBLICACION"] for r in cruce_clean if r["NIVEL"] == "PREGRADO")
    inst_pre = sum(r["INSTITUCIONAL_VIGENTE"] for r in cruce_clean if r["NIVEL"] == "PREGRADO")
    pub_pos = sum(r["PUBLICACION"] for r in cruce_clean if r["NIVEL"] == "POSGRADO_POSTITULO")
    inst_pos = sum(r["INSTITUCIONAL_VIGENTE"] for r in cruce_clean if r["NIVEL"] == "POSGRADO_POSTITULO")
    pub_total = sum(r["PUBLICACION"] for r in cruce_clean)
    inst_total = sum(r["INSTITUCIONAL_VIGENTE"] for r in cruce_clean)
    diff_total = pub_total - inst_total
    pos_sum = sum(r["DIFERENCIA"] for r in positives)
    neg_sum = sum(r["DIFERENCIA"] for r in negatives)

    # Inventario de CSV históricos ya auditados.
    inv_csv_path = CSV_AUDIT_DIR / "01_INVENTARIO_ARCHIVOS.csv"
    universe_csv_path = CSV_AUDIT_DIR / "03_COMPARACION_UNIVERSO.csv"
    csv_inventory = []
    with inv_csv_path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        csv_inventory = list(reader)
    universe_rows = {}
    with universe_csv_path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            universe_rows[row["archivo"]] = row

    git_branch = get_git(["branch", "--show-current"])
    git_commit = get_git(["rev-parse", "HEAD"])

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "01_RESUMEN_GERENCIAL"
    ws_informed = wb.create_sheet("02_INFORMADO_A_SIES")
    ws_gap = wb.create_sheet("03_BRECHA_POR_CARRERA")
    ws_trace = wb.create_sheet("04_CONTROL_TRAZABILIDAD")

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    # H1 Resumen gerencial
    style_title(ws_summary, "A1", "REPORTE GERENCIAL DE BRECHAS - MATRÍCULA UNIFICADA SIES 2026")
    ws_summary["A2"] = "La diferencia neta de 226 se encuentra identificada y distribuida por código de carrera. No corresponde a una nómina de 226 RUT, porque la publicación no contiene identificadores individuales."
    ws_summary["A2"].font = Font(bold=True, color="9C5700")
    ws_summary["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws_summary["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws_summary.merge_cells("A2:P3")

    cards = [
        ("Publicación Pregrado", pub_pre, "D9EAF7"),
        ("Institucional Pregrado vigente", inst_pre, "D9EAF7"),
        ("Diferencia Pregrado", signed(pub_pre - inst_pre), "FCE4D6"),
        ("Publicación Posgrado/Postítulo", pub_pos, "E2F0D9"),
        ("Institucional Posgrado/Postítulo", inst_pos, "E2F0D9"),
        ("Diferencia Posgrado/Postítulo", signed(pub_pos - inst_pos), "E2F0D9"),
        ("Publicación total", pub_total, "D9EAF7"),
        ("Institucional total", inst_total, "D9EAF7"),
        ("Diferencia neta", signed(diff_total), "FCE4D6"),
        ("Códigos involucrados con diferencia", len(diff_rows), "FCE4D6"),
        ("Códigos diferencia positiva", len(positives), "FCE4D6"),
        ("Códigos diferencia negativa", len(negatives), "F4CCCC"),
        ("Códigos sin diferencia", len(iguales), "E2F0D9"),
        ("Registros institucionales VIG=0", total_vig["0"], "E7E6E6"),
    ]
    start_cols = [1, 3, 5, 7, 9, 11, 13]
    for idx, (title, value, fill) in enumerate(cards):
        row = 5 + (idx // 7) * 3
        col = start_cols[idx % 7]
        add_card(ws_summary, row, col, title, value, fill)

    summary_headers = [
        "ORDEN_BRECHA",
        "CODIGO_SIES",
        "NOMBRE_CARRERA",
        "NIVEL",
        "PUBLICACION",
        "INSTITUCIONAL_VIGENTE",
        "DIFERENCIA",
        "TIPO_BRECHA",
        "RUT_UNICOS_INSTITUCIONALES_VIGENTES",
        "HOMBRES_INSTITUCIONALES_VIGENTES",
        "MUJERES_INSTITUCIONALES_VIGENTES",
        "OTRAS_CATEGORIAS_SEXO",
        "SIN_INFORMACION_SEXO",
        "REGISTROS_VIG_0",
        "PORCENTAJE_SOBRE_BRECHA_POSITIVA_BRUTA",
        "OBSERVACION",
    ]
    summary_rows = []
    for r in diff_rows:
        pct = r["DIFERENCIA"] / pos_sum if r["DIFERENCIA"] > 0 and pos_sum else None
        summary_rows.append([
            r["ORDEN_BRECHA"],
            r["CODIGO_SIES"],
            r["NOMBRE_CARRERA"],
            r["NIVEL"],
            r["PUBLICACION"],
            r["INSTITUCIONAL_VIGENTE"],
            r["DIFERENCIA"],
            r["TIPO_BRECHA"],
            r["RUT_UNICOS_INSTITUCIONALES_VIGENTES"],
            r["HOMBRES_INSTITUCIONALES_VIGENTES"],
            r["MUJERES_INSTITUCIONALES_VIGENTES"],
            r["OTRAS_CATEGORIAS_SEXO"],
            r["SIN_INFORMACION_SEXO"],
            r["REGISTROS_VIG_0"],
            pct if pct is not None else "No aplica",
            r["OBSERVACION"],
        ])
    table_end, _ = write_table(ws_summary, 13, 1, summary_headers, summary_rows, "tblResumenBrechasGerencial", "TableStyleMedium2")
    ws_summary.freeze_panes = "A14"
    for row in range(14, table_end + 1):
        diff_cell = ws_summary.cell(row, 7)
        diff_cell.number_format = '+#,##0;-#,##0;0'
        if isinstance(diff_cell.value, int) and diff_cell.value > 0:
            diff_cell.fill = PatternFill("solid", fgColor="FCE4D6")
        elif isinstance(diff_cell.value, int) and diff_cell.value < 0:
            diff_cell.fill = PatternFill("solid", fgColor="F4CCCC")
        ws_summary.cell(row, 15).number_format = "0.0%"

    chart_top = positives[:10]
    chart_start = table_end + 3
    ws_summary.cell(chart_start, 1, "Top 10 diferencias positivas")
    ws_summary.cell(chart_start, 1).font = Font(bold=True, color="17365D")
    ws_summary.cell(chart_start + 1, 1, "CODIGO_SIES")
    ws_summary.cell(chart_start + 1, 2, "DIFERENCIA")
    for i, r in enumerate(chart_top, chart_start + 2):
        ws_summary.cell(i, 1, r["CODIGO_SIES"])
        ws_summary.cell(i, 2, r["DIFERENCIA"])
    chart = BarChart()
    chart.title = "Diez mayores diferencias positivas"
    chart.y_axis.title = "Matrículas"
    chart.x_axis.title = "Código SIES"
    data = Reference(ws_summary, min_col=2, min_row=chart_start + 1, max_row=chart_start + 1 + len(chart_top))
    cats = Reference(ws_summary, min_col=1, min_row=chart_start + 2, max_row=chart_start + 1 + len(chart_top))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 7
    chart.width = 18
    ws_summary.add_chart(chart, "D" + str(chart_start))

    # H2 Informado a SIES
    for col, header in enumerate(pre_headers, 1):
        ws_informed.cell(1, col, header)
    for row_idx, r in enumerate(pre_rows, 2):
        for col_idx, h in enumerate(pre_headers, 1):
            ws_informed.cell(row_idx, col_idx, r.get(h))
    raw_ref = f"A1:{get_column_letter(len(pre_headers))}{len(pre_rows) + 1}"
    tab = Table(displayName="tblInformadoSIES", ref=raw_ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False,
                                        showRowStripes=True, showColumnStripes=False)
    ws_informed.add_table(tab)
    ws_informed.freeze_panes = "A2"
    ws_informed.auto_filter.ref = raw_ref

    # H3 Brecha por carrera
    style_title(ws_gap, "A1", "BRECHA POR CARRERA")
    ws_gap["A2"] = "La brecha está identificada por código y cantidad, pero no por persona. La publicación no contiene RUT ni sexo."
    ws_gap["A2"].font = Font(bold=True, color="9C5700")
    ws_gap["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws_gap.merge_cells("A2:Q3")
    reconciliation = [
        ["46 códigos con publicación mayor", len(positives), pos_sum],
        ["2 códigos con institucional mayor", len(negatives), neg_sum],
        ["Diferencia neta", "", diff_total],
        ["18 códigos sin diferencia", len(iguales), 0],
        ["2 códigos solo en publicación", len(solo_publicacion), sum(r["DIFERENCIA"] for r in solo_publicacion)],
        ["0 códigos solo en institucional", len(solo_institucional), 0],
    ]
    write_table(ws_gap, 5, 1, ["CONTROL", "CANTIDAD", "VALOR"], reconciliation, "tblConciliacionBrecha", "TableStyleMedium4")
    gap_headers = [
        "CODIGO_SIES",
        "NOMBRE_CARRERA",
        "NIVEL",
        "PUBLICACION",
        "INSTITUCIONAL_VIGENTE",
        "DIFERENCIA",
        "SIGNO",
        "TIPO_BRECHA",
        "SOLO_PUBLICACION",
        "RUT_UNICOS_INSTITUCIONALES_VIGENTES",
        "REGISTROS_VIG_0",
        "HOMBRES_INSTITUCIONALES_VIGENTES",
        "MUJERES_INSTITUCIONALES_VIGENTES",
        "OTRAS_CATEGORIAS_SEXO",
        "SIN_INFORMACION_SEXO",
        "FUENTE_NOMBRE_CARRERA",
        "OBSERVACION",
    ]
    gap_rows = []
    for r in diff_rows:
        gap_rows.append([
            r["CODIGO_SIES"],
            r["NOMBRE_CARRERA"],
            r["NIVEL"],
            r["PUBLICACION"],
            r["INSTITUCIONAL_VIGENTE"],
            r["DIFERENCIA"],
            "POSITIVA" if r["DIFERENCIA"] > 0 else "NEGATIVA",
            r["TIPO_BRECHA"],
            "SI" if r["TIPO_BRECHA"] == "SOLO_PUBLICACION" else "NO",
            r["RUT_UNICOS_INSTITUCIONALES_VIGENTES"],
            r["REGISTROS_VIG_0"],
            r["HOMBRES_INSTITUCIONALES_VIGENTES"],
            r["MUJERES_INSTITUCIONALES_VIGENTES"],
            r["OTRAS_CATEGORIAS_SEXO"],
            r["SIN_INFORMACION_SEXO"],
            r["FUENTE_NOMBRE_CARRERA"],
            r["OBSERVACION"],
        ])
    gap_end, _ = write_table(ws_gap, 13, 1, gap_headers, gap_rows, "tblBrechaPorCarrera", "TableStyleMedium2")
    ws_gap.freeze_panes = "A14"
    for row in range(14, gap_end + 1):
        ws_gap.cell(row, 6).number_format = '+#,##0;-#,##0;0'
        if ws_gap.cell(row, 6).value > 0:
            ws_gap.cell(row, 6).fill = PatternFill("solid", fgColor="FCE4D6")
        else:
            ws_gap.cell(row, 6).fill = PatternFill("solid", fgColor="F4CCCC")

    # H4 Control trazabilidad
    style_title(ws_trace, "A1", "CONTROL Y TRAZABILIDAD")
    row = 3
    style_section(ws_trace, row, "BLOQUE A — IDENTIFICACIÓN", 1, 5)
    row += 1
    ident = [
        ["Proceso", "Matrícula Unificada SIES 2026"],
        ["Subproyecto", "Auditoría gerencial de diferencias entre publicación e información institucional de Pregrado y Posgrado/Postítulo"],
        ["Año del proceso", 2026],
        ["Año de referencia", 2026],
        ["Repositorio", str(REPO)],
        ["Fecha y hora de ejecución", datetime.now().isoformat(timespec="seconds")],
        ["Rama Git", git_branch],
        ["Commit Git", git_commit],
        ["Script o comando utilizado", str(Path(__file__).resolve())],
        ["Archivo de salida", str(xlsx_path)],
        ["Originales modificados", "NO"],
    ]
    end, _ = write_table(ws_trace, row, 1, ["CAMPO", "VALOR"], ident, "tblIdentificacion", "TableStyleMedium2")
    row = end + 3

    style_section(ws_trace, row, "BLOQUE B — FUENTES", 1, 8)
    row += 1
    fuentes = [
        ["Libro reconstruido", str(SOURCE_BOOK), SOURCE_BOOK.name, "", sha256(SOURCE_BOOK), PREGRADO_SHEET, "Fuente gobernante de comparación"],
        ["Auditoría seis CSV", str(CSV_AUDIT_DIR), CSV_AUDIT_DIR.name, "", "", "", "Carpeta de auditoría histórica"],
        ["Manual / fuente funcional VIG", str(REPO / "Manual_Matrícula_Unificada_2026.pdf"), "Manual_Matrícula_Unificada_2026.pdf", "", sha256(REPO / "Manual_Matrícula_Unificada_2026.pdf") if (REPO / "Manual_Matrícula_Unificada_2026.pdf").exists() else "", "", "Fuente funcional disponible en proyecto"],
    ]
    for c in csv_inventory:
        u = universe_rows.get(c["Nombre archivo"], {})
        fuentes.append(["CSV histórico", c.get("Ruta"), c.get("Nombre archivo"), c.get("Cantidad filas físicas"), c.get("Hash SHA-256"), "", u.get("clasificacion_final", "")])
    end, _ = write_table(ws_trace, row, 1, ["TIPO", "RUTA", "NOMBRE", "FILAS", "SHA256", "HOJA", "EVIDENCIA / CLASIFICACION"], fuentes, "tblFuentes", "TableStyleMedium2")
    row = end + 3

    style_section(ws_trace, row, "BLOQUE C — METODOLOGÍA", 1, 6)
    row += 1
    metodo = [
        ["Publicación", "Fuente agregada por código; no contiene RUT ni sexo individual."],
        ["Institucional", "Fuente con detalle individual en 08_PREGRADO_COMPLETO."],
        ["Unidad de conciliación", "Código SIES."],
        ["Fórmula", "Publicación menos institucional vigente."],
        ["Vigencia comparable", "VIG=1 + VIG=2."],
        ["VIG=0", "Registro informado, pero no matrícula vigente comparable."],
        ["RUT único", "Conteo distinto dentro de cada código; no se suma como total de personas."],
        ["Sexo", "Composición de registros institucionales, no de la brecha."],
        ["Columnas documentales", "RUT/documento: TIPO_DOC + N_DOC + DV; CODCLI: no existe en 08_PREGRADO_COMPLETO; sexo: SEXO; código carrera: COD_CAR; vigencia: VIG; sede: COD_SED; jornada: JOR; versión: VERSION."],
        ["Mapeo sexo", "H = hombres; M = mujeres; otros valores se conservan en OTRAS_CATEGORIAS_SEXO; vacío en SIN_INFORMACION_SEXO."],
    ]
    end, _ = write_table(ws_trace, row, 1, ["CAMPO", "METODO"], metodo, "tblMetodologia", "TableStyleMedium2")
    row = end + 3

    style_section(ws_trace, row, "BLOQUE D — LIMITACIONES", 1, 6)
    row += 1
    limitations = [
        ["No existe una nómina de 226 RUT asociada a la diferencia."],
        ["La publicación no contiene identificadores individuales."],
        ["La brecha solo puede localizarse por código de carrera y cantidad."],
        ["No puede atribuirse sexo a los registros de diferencia."],
        ["Los seis CSV históricos son fuentes parciales anteriores y ninguno reproduce el consolidado final."],
        ["La suma de RUT únicos por carrera no equivale necesariamente al total de personas únicas institucionales."],
    ]
    end, _ = write_table(ws_trace, row, 1, ["LIMITACION"], limitations, "tblLimitaciones", "TableStyleMedium2")
    row = end + 3

    style_section(ws_trace, row, "BLOQUE E — VALIDACIONES", 1, 7)
    validation_start_row = row + 1
    validation_headers = ["CONTROL", "ESPERADO", "OBTENIDO", "ESTADO", "EVIDENCIA"]
    validations = []
    def add_val(control: str, esperado: Any, obtenido: Any, evidence: str, ok: bool | None = None):
        if ok is None:
            ok = str(esperado) == str(obtenido)
        validations.append([control, esperado, obtenido, "OK" if ok else "REVISAR", evidence])

    add_val("Cuatro hojas exactas", "01_RESUMEN_GERENCIAL;02_INFORMADO_A_SIES;03_BRECHA_POR_CARRERA;04_CONTROL_TRAZABILIDAD", "PENDIENTE_POST_GUARDADO", "Validación post guardado", False)
    add_val("4.105 registros en 02_INFORMADO_A_SIES", 4105, len(pre_rows), "08_PREGRADO_COMPLETO primeras 32 columnas")
    add_val("32 columnas", 32, len(pre_headers), "Estructura oficial pregrado")
    add_val("VIG=0 igual a 960", 960, total_vig["0"], "Conteo columna VIG")
    add_val("VIG=1 igual a 3.145", 3145, total_vig["1"], "Conteo columna VIG")
    add_val("VIG=2 igual a 0", 0, total_vig["2"], "Conteo columna VIG")
    add_val("Publicación Pregrado igual a 3.371", 3371, pub_pre, CRUCE_SHEET)
    add_val("Institucional Pregrado igual a 3.145", 3145, inst_pre, CRUCE_SHEET)
    add_val("Diferencia Pregrado igual a +226", 226, pub_pre - inst_pre, CRUCE_SHEET)
    add_val("Posgrado/Postítulo igual a 54 vs. 54", "54 vs 54", f"{pub_pos} vs {inst_pos}", CRUCE_SHEET)
    add_val("66 códigos totales", 66, len(cruce_clean), CRUCE_SHEET)
    add_val("46 diferencias positivas, suma +228", "46 / 228", f"{len(positives)} / {pos_sum}", CRUCE_SHEET)
    add_val("2 diferencias negativas, suma -2", "2 / -2", f"{len(negatives)} / {neg_sum}", CRUCE_SHEET)
    add_val("18 códigos iguales", 18, len(iguales), CRUCE_SHEET)
    add_val("2 códigos solo publicación", 2, len(solo_publicacion), CRUCE_SHEET)
    add_val("0 códigos solo institucional", 0, len(solo_institucional), CRUCE_SHEET)
    add_val("Diferencia neta +226", 226, diff_total, CRUCE_SHEET)
    add_val("Fórmulas con error", 0, "PENDIENTE_POST_GUARDADO", "openpyxl scan", False)
    add_val("Hipervínculos externos", 0, "PENDIENTE_POST_GUARDADO", "OOXML rels", False)
    add_val("Relaciones externas", 0, "PENDIENTE_POST_GUARDADO", "OOXML rels", False)
    add_val("Errores XML", 0, "PENDIENTE_POST_GUARDADO", "zipfile/XML parse", False)
    add_val("Originales modificados", "NO", "PENDIENTE_POST_GUARDADO", "Hash before/after", False)

    val_end, _ = write_table(ws_trace, validation_start_row, 1, validation_headers, validations, "tblValidaciones", "TableStyleMedium2")
    ws_trace.freeze_panes = "A2"

    # Formatting common
    for ws in [ws_summary, ws_informed, ws_gap, ws_trace]:
        ws.sheet_view.showGridLines = False
        auto_width(ws)
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                cell.font = copy(cell.font)
                if cell.row == 1:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = ws.freeze_panes or "A2"

    # Specific widths
    ws_summary.column_dimensions["C"].width = 44
    ws_gap.column_dimensions["B"].width = 44
    ws_informed.column_dimensions["F"].width = 24
    ws_informed.column_dimensions["A"].width = 10
    ws_informed.column_dimensions["B"].width = 13
    ws_informed.column_dimensions["C"].width = 8
    for ws in [ws_summary, ws_gap]:
        for col in ["E", "F", "G", "I", "J", "K", "M", "N"]:
            if col in ws.column_dimensions:
                ws.column_dimensions[col].width = 14
    for ws in [ws_summary, ws_gap]:
        for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row_cells:
                if isinstance(cell.value, int):
                    cell.number_format = '#,##0'
    for ws in [ws_summary, ws_gap]:
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.value in {"OK", "REPORTE_GERENCIAL_MU2026_CONSTRUIDO_Y_VALIDADO"}:
                    cell.fill = PatternFill("solid", fgColor="E2F0D9")
                elif cell.value == "REVISAR":
                    cell.fill = PatternFill("solid", fgColor="FCE4D6")

    # Save first pass.
    wb.save(xlsx_path)
    wb.close()

    # Post-save validation.
    read_wb = load_workbook(xlsx_path, read_only=True, data_only=False)
    sheet_names = read_wb.sheetnames
    informed_ws = read_wb["02_INFORMADO_A_SIES"]
    post_vig = Counter()
    headers_post = [c.value for c in next(informed_ws.iter_rows(min_row=1, max_row=1))]
    vig_idx = headers_post.index("VIG") + 1
    for r in informed_ws.iter_rows(min_row=2, max_row=informed_ws.max_row, values_only=True):
        post_vig[norm_text(r[vig_idx - 1])] += 1
    formula_errors = validate_formula_errors(xlsx_path)
    xml_errors, external_rels, external_hyperlinks = parse_xml_and_rels(xlsx_path)
    source_hashes_after = {str(p): sha256(p) for p in expected_sources if p.exists()}
    originals_unchanged = source_hashes_before == source_hashes_after

    post_checks = {
        "sheets_exact": sheet_names == ["01_RESUMEN_GERENCIAL", "02_INFORMADO_A_SIES", "03_BRECHA_POR_CARRERA", "04_CONTROL_TRAZABILIDAD"],
        "raw_rows": informed_ws.max_row - 1 == 4105,
        "raw_cols": informed_ws.max_column == 32,
        "vig0": post_vig["0"] == 960,
        "vig1": post_vig["1"] == 3145,
        "vig2": post_vig["2"] == 0,
        "formula_errors": formula_errors == 0,
        "xml_errors": xml_errors == 0,
        "external_rels": external_rels == 0,
        "external_hyperlinks": external_hyperlinks == 0,
        "originals_unchanged": originals_unchanged,
    }
    read_wb.close()

    # Rewrite validation statuses into workbook.
    wb2 = load_workbook(xlsx_path)
    ws2 = wb2["04_CONTROL_TRAZABILIDAD"]
    # Table validation starts at validation_start_row; data rows start +1.
    for r in range(validation_start_row + 1, ws2.max_row + 1):
        control = ws2.cell(r, 1).value
        if control == "Cuatro hojas exactas":
            ws2.cell(r, 3, ";".join(sheet_names))
            ws2.cell(r, 4, "OK" if post_checks["sheets_exact"] else "REVISAR")
        elif control == "Fórmulas con error":
            ws2.cell(r, 3, formula_errors)
            ws2.cell(r, 4, "OK" if formula_errors == 0 else "REVISAR")
        elif control == "Hipervínculos externos":
            ws2.cell(r, 3, external_hyperlinks)
            ws2.cell(r, 4, "OK" if external_hyperlinks == 0 else "REVISAR")
        elif control == "Relaciones externas":
            ws2.cell(r, 3, external_rels)
            ws2.cell(r, 4, "OK" if external_rels == 0 else "REVISAR")
        elif control == "Errores XML":
            ws2.cell(r, 3, xml_errors)
            ws2.cell(r, 4, "OK" if xml_errors == 0 else "REVISAR")
        elif control == "Originales modificados":
            ws2.cell(r, 3, "NO" if originals_unchanged else "SI")
            ws2.cell(r, 4, "OK" if originals_unchanged else "REVISAR")
    for row_cells in ws2.iter_rows():
        for cell in row_cells:
            if cell.value == "OK":
                cell.fill = PatternFill("solid", fgColor="E2F0D9")
            elif cell.value == "REVISAR":
                cell.fill = PatternFill("solid", fgColor="FCE4D6")
    wb2.save(xlsx_path)
    wb2.close()

    # Revalidate after rewrite.
    formula_errors = validate_formula_errors(xlsx_path)
    xml_errors, external_rels, external_hyperlinks = parse_xml_and_rels(xlsx_path)
    read_wb = load_workbook(xlsx_path, read_only=True, data_only=False)
    final_sheets = read_wb.sheetnames
    final_rows = read_wb["02_INFORMADO_A_SIES"].max_row - 1
    final_cols = read_wb["02_INFORMADO_A_SIES"].max_column
    read_wb.close()
    all_ok = (
        final_sheets == ["01_RESUMEN_GERENCIAL", "02_INFORMADO_A_SIES", "03_BRECHA_POR_CARRERA", "04_CONTROL_TRAZABILIDAD"]
        and final_rows == 4105
        and final_cols == 32
        and post_vig["0"] == 960
        and post_vig["1"] == 3145
        and post_vig["2"] == 0
        and pub_pre == 3371
        and inst_pre == 3145
        and pub_pos == 54
        and inst_pos == 54
        and pub_total == 3425
        and inst_total == 3199
        and diff_total == 226
        and len(positives) == 46
        and pos_sum == 228
        and len(negatives) == 2
        and neg_sum == -2
        and len(iguales) == 18
        and len(solo_publicacion) == 2
        and len(solo_institucional) == 0
        and formula_errors == 0
        and xml_errors == 0
        and external_rels == 0
        and external_hyperlinks == 0
        and originals_unchanged
    )
    estado = "REPORTE_GERENCIAL_MU2026_CONSTRUIDO_Y_VALIDADO" if all_ok else "REPORTE_GERENCIAL_MU2026_NO_LIBERADO"

    repo_hash = sha256(xlsx_path)
    desktop_hash = None
    desktop_backup = None
    if all_ok:
        if DESKTOP_FINAL.exists():
            desktop_backup = DESKTOP_FINAL.with_name(f"{DESKTOP_FINAL.stem}_RESPALDO_{ts}{DESKTOP_FINAL.suffix}")
            shutil.copy2(DESKTOP_FINAL, desktop_backup)
        shutil.copy2(xlsx_path, DESKTOP_FINAL)
        desktop_hash = sha256(DESKTOP_FINAL)
        if desktop_hash != repo_hash:
            estado = "REPORTE_GERENCIAL_MU2026_NO_LIBERADO"

    # Inventario SHA.
    inventory_rows = []
    inventory_sources = [SOURCE_BOOK, inv_csv_path, universe_csv_path, CSV_AUDIT_DIR / "04_COMPARACION_CODIGOS.csv", xlsx_path]
    for c in csv_inventory:
        ruta = Path(c["Ruta"])
        if ruta.exists():
            inventory_sources.append(ruta)
    if (REPO / "Manual_Matrícula_Unificada_2026.pdf").exists():
        inventory_sources.append(REPO / "Manual_Matrícula_Unificada_2026.pdf")
    seen = set()
    for p in inventory_sources:
        if str(p) in seen or not p.exists():
            continue
        seen.add(str(p))
        fi = file_info(p)
        inventory_rows.append({
            "tipo": "SALIDA" if p == xlsx_path else "FUENTE",
            **fi,
        })
    with inventory_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["tipo", "ruta", "nombre", "tamano_bytes", "fecha_modificacion", "sha256"])
        writer.writeheader()
        writer.writerows(inventory_rows)

    manifest = {
        "estado": estado,
        "proceso": "Matrícula Unificada SIES 2026",
        "subproyecto": "Auditoría gerencial de diferencias entre publicación e información institucional de Pregrado y Posgrado/Postítulo",
        "anio_proceso": 2026,
        "anio_referencia": 2026,
        "fecha_ejecucion": datetime.now().isoformat(timespec="seconds"),
        "repositorio": str(REPO),
        "rama_git": git_branch,
        "commit_git": git_commit,
        "fuente_gobernante": file_info(SOURCE_BOOK),
        "hojas_fuente": {
            "publicacion": PUBLICACION_SHEET,
            "pregrado": PREGRADO_SHEET,
            "posgrado": POSGRADO_SHEET,
            "cruce": CRUCE_SHEET,
        },
        "totales": {
            "publicacion_pregrado": pub_pre,
            "institucional_pregrado": inst_pre,
            "diferencia_pregrado": pub_pre - inst_pre,
            "publicacion_posgrado_postitulo": pub_pos,
            "institucional_posgrado_postitulo": inst_pos,
            "diferencia_posgrado_postitulo": pub_pos - inst_pos,
            "publicacion_total": pub_total,
            "institucional_total": inst_total,
            "diferencia_neta": diff_total,
            "vig0": total_vig["0"],
            "vig1": total_vig["1"],
            "vig2": total_vig["2"],
            "codigos_totales": len(cruce_clean),
            "codigos_sin_diferencia": len(iguales),
            "codigos_diferencia_positiva": len(positives),
            "suma_diferencias_positivas": pos_sum,
            "codigos_diferencia_negativa": len(negatives),
            "suma_diferencias_negativas": neg_sum,
            "codigos_solo_publicacion": [r["CODIGO_SIES"] for r in solo_publicacion],
            "codigos_solo_institucional": [r["CODIGO_SIES"] for r in solo_institucional],
        },
        "validaciones": {
            "hojas": final_sheets,
            "filas_informado_sies": final_rows,
            "columnas_informado_sies": final_cols,
            "formula_errors": formula_errors,
            "xml_errors": xml_errors,
            "external_rels": external_rels,
            "external_hyperlinks": external_hyperlinks,
            "originales_modificados": "NO" if originals_unchanged else "SI",
            "desktop_hash_igual": desktop_hash == repo_hash if desktop_hash else False,
        },
        "archivos_generados": {
            "excel_repositorio": str(xlsx_path),
            "excel_escritorio": str(DESKTOP_FINAL) if desktop_hash else None,
            "desktop_backup_previo": str(desktop_backup) if desktop_backup else None,
            "validacion_txt": str(validation_path),
            "manifest": str(manifest_path),
            "inventario_fuentes_sha256": str(inventory_path),
            "sha256_excel_repositorio": repo_hash,
            "sha256_excel_escritorio": desktop_hash,
        },
        "limitaciones": [
            "No existe una nómina de 226 RUT asociada a la diferencia.",
            "La publicación no contiene identificadores individuales.",
            "La brecha solo puede localizarse por código de carrera y cantidad.",
            "No puede atribuirse sexo a los registros de diferencia.",
            "Los seis CSV históricos son fuentes parciales anteriores y ninguno reproduce el consolidado final.",
            "La suma de RUT únicos por carrera no equivale necesariamente al total de personas únicas institucionales.",
        ],
    }
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")

    validation_text = f"""Estado: {estado}
Ruta Excel repositorio: {xlsx_path}
Ruta Excel Escritorio: {DESKTOP_FINAL if desktop_hash else 'NO_COPIADO'}
SHA-256 repositorio: {repo_hash}
SHA-256 Escritorio: {desktop_hash or 'NO_COPIADO'}
Cantidad de hojas: {len(final_sheets)}
Hojas: {'; '.join(final_sheets)}
Filas de 02_INFORMADO_A_SIES: {final_rows}
Cantidad de columnas: {final_cols}
Conteos VIG=0, VIG=1 y VIG=2: {post_vig['0']}, {post_vig['1']}, {post_vig['2']}
Publicación total: {pub_total}
Institucional total: {inst_total}
Diferencia neta: {diff_total}
Cantidad y suma de diferencias positivas: {len(positives)} / {pos_sum}
Cantidad y suma de diferencias negativas: {len(negatives)} / {neg_sum}
Códigos solo publicación: {', '.join(r['CODIGO_SIES'] for r in solo_publicacion)}
Errores XML: {xml_errors}
Relaciones externas: {external_rels}
Hipervínculos externos: {external_hyperlinks}
Fórmulas con error: {formula_errors}
Originales modificados: {'NO' if originals_unchanged else 'SI'}
Conclusión gerencial: La diferencia neta de 226 matrículas se encuentra completamente distribuida y cuantificada por código de carrera. Sin embargo, la publicación no contiene identificadores individuales, por lo que no es posible reconstruir una nómina de 226 RUT ni determinar la composición por sexo de esa brecha. Los conteos de RUT y sexo incorporados corresponden exclusivamente a los registros institucionales informados.
"""
    validation_path.write_text(validation_text, encoding="utf-8")

    print(validation_text)
    print(f"Manifest: {manifest_path}")
    print(f"Inventario: {inventory_path}")
    return 0 if all_ok and desktop_hash == repo_hash else 1


if __name__ == "__main__":
    raise SystemExit(main())
