"""Construye documentación agregada UMAS/VIG; no asigna VIG a estudiantes.

Lee exclusivamente pares de estado/situación y metadatos. Los escenarios de
decisión son sintéticos. No ejecuta motores congelados ni cambia indicadores.
"""
from pathlib import Path
from collections import Counter
from datetime import datetime
from zoneinfo import ZoneInfo
import ast
import csv
import hashlib
import json
import subprocess

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/reglas_negocio/20261020_estados_umas'
RAW = Path('/Users/alexi/Downloads/DATOSDEALUMNOS_12112.xlsx')
MU = ROOT / 'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
BOOK = ROOT / 'outputs/retencion_intermedia_t2_t3/20261020_221157/CONTINUIDAD_IPSS_INGRESOS_2025_T2_T3_EN_MU2026_20261020_221157.xlsx'
GATE = ROOT / 'requerimiento_bettersoft_uplus_2026'


def sha(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')


def write_tsv(path, rows):
    with path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]), delimiter='\t', lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def decide(evidence):
    """C: evaluador de escenarios sintéticos, no adaptador a DatosAlumnos.

    Cada SI/NO implica evidencia independiente revisada, no inferencia de un
    estado, situación, fecha aislada ni ausencia de registro.
    """
    if evidence['contexto'] != 'SI' or evidence['contradiccion'] != 'NO':
        return 'RE', None
    if evidence['matricula_vigente'] == 'NO':
        return 'R0', 0
    if evidence['matricula_vigente'] != 'SI':
        return 'RE', None
    if evidence['condicion_terminal'] == 'SI':
        return 'R2', 2
    if evidence['condicion_terminal'] == 'NO':
        return 'R1', 1
    return 'RE', None


def main():
    before = {'rama': subprocess.check_output(['git', 'branch', '--show-current'], cwd=ROOT, text=True).strip(),
              'HEAD': subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()}
    sources = [
        ('A', ROOT / 'Manual_Matrícula_Unificada_2026.pdf', 'p.7 §2.vi; p.10 suspensiones; p.31 Anexo7 Cuadro1 AF'),
        ('D/B', GATE / 'REPORTE_CIERRE_GATE09_UPLUS_SIES_2026.md', '§C C09-01; §E1–3'),
        ('D/B', GATE / 'MANIFIESTO_CONGELACION_v1.0.md', 'C09_01 CONFIRMADA'),
        ('D/B', GATE / 'resultados/DICCIONARIO_MAESTRO_UPLUS_SIES_2026_v1.0.xlsx', '12_GATE09_RESOLUCION fila2;18_CONFIRMACION_GATE10 fila2;05_DICCIONARIO_BETTERSOFT UP39/40/42/43/59/60/66/68;06_CATALOGOS_HOMOLOGACIONES J79:K79'),
        ('C', ROOT / 'gobernanza_catalogos/gob_datosalumnos_estadoacademico_situacion.tsv', '29 parejas; VIG_ESPERADO es antecedente descartado como regla por Gate10'),
        ('C', ROOT / 'procesos/caracterizacion/caracterizacion.py', '_MAPA_ESTADO_OFICIAL; consulta AST sin ejecutar módulo'),
        ('B', RAW, 'Hoja1; ESTADOACADEMICO y SITUACION; solo conteos agregados'),
    ]
    inventory = [{'clasificacion': cl, 'ruta': str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p),
                  'sha256': sha(p), 'localizador': loc, 'fuente_original_modificada': False} for cl, p, loc in sources]
    assert sha(RAW) == 'fdd47d99b934479f39d8db6d00eceb40714d3b8c7baa1e3a89a9288b2b8bc471'
    preserved = {str(p.relative_to(ROOT)): sha(p) for p in [MU, BOOK]}
    tree = ast.parse((ROOT / 'procesos/caracterizacion/caracterizacion.py').read_text())
    mapping = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                   and any(isinstance(t, ast.Name) and t.id == '_MAPA_ESTADO_OFICIAL' for t in n.targets))
    with (ROOT / 'gobernanza_catalogos/gob_datosalumnos_estadoacademico_situacion.tsv').open() as f:
        old_rows = list(csv.DictReader(f, delimiter='\t'))
    old = {(r['ESTADOACADEMICO'].strip(), r['SITUACION'].strip()): r['VIG_ESPERADO'] for r in old_rows}
    wb = load_workbook(RAW, read_only=True, data_only=True)
    ws = wb['Hoja1']
    data = ws.iter_rows(values_only=True)
    headers = next(data)
    ei, si = headers.index('ESTADOACADEMICO'), headers.index('SITUACION')
    counts = Counter((str(r[ei]), str(r[si])) for r in data if any(v is not None for v in r))
    wb.close()
    assert len(headers) == 67 and sum(counts.values()) == 15136 and len(counts) == 32
    rows = []
    special = {
        'EGRESADO': 'Verificar fecha de egreso, matrícula al corte y actividades terminales. EGRESADO no basta para R2.',
        'TITULADO': 'Verificar fecha del título y trayectoria al corte. Título ya obtenido no acredita R2 para ese título; evaluar otra matrícula separadamente. No asignar 0/2 por etiqueta.',
        'SUSPENDIDO': 'Manual p.10: suspensión admite 0/1/2 conforme a reglamento y matrícula. Falta reglamento aplicable para resolver desde la etiqueta.',
        'ELIMINADO': 'Revisar fecha y alcance del evento; no trasladar eliminación actual a una matrícula histórica ni inferir matrícula de destino.',
        'VIGENTE': 'Estado actual no prueba matrícula vigente del período/corte ni distingue por sí solo R1/R2.',
    }
    for (state, situation), n in sorted(counts.items()):
        code = situation.split(' - ', 1)[0].strip()
        classification = mapping.get((code, state), 'POR_REVISAR')
        note = special[state]
        if code == '2':
            note = special['SUSPENDIDO']
        if code == '16':
            note += ' Significado de situación16 sin catálogo confirmado; no inferir causa disciplinaria individual.'
        rows.append({
            'ESTADOACADEMICO_ORIGINAL': state, 'SITUACION_ORIGINAL': situation,
            'CODIGO_SITUACION_TECNICO': code, 'SITUACION_BUSQUEDA_TRIM': situation.strip(),
            'REGISTROS_OBSERVADOS_B': n, 'CLASIFICACION_INTERNA_C': classification,
            'RESPALDO_PAREJA': 'C: mapa integrado _MAPA_ESTADO_OFICIAL' if classification != 'POR_REVISAR' else 'B: etiqueta; E: definición pendiente',
            'VIG_HISTORICO_DESCARTADO_NO_USAR': old.get((state, situation.strip()), ''),
            'VIG_1_CONDICION_C': 'R1: misma matrícula/período/corte; matrícula vigente y condición no terminal acreditadas',
            'VIG_2_CONDICION_C': 'R2: misma matrícula/período/corte; matrícula vigente, plan finalizado, actividades terminales en curso y título no obtenido al corte',
            'VIG_0_CONDICION_C': 'R0: condición sin matrícula acreditada al corte; no desde ausencia de evidencia',
            'SIN_EVIDENCIA_E': 'RE: pendiente; VIG vacío; conservar VIG informado en columna separada',
            'VIG_ASIGNADO_A_ESTUDIANTES': '', 'OBSERVACION_C_E': note,
            'FUENTE_DEFINICION_A': 'Manual MU2026 p.7 §2.vi;p.31 AF;p.10 para suspensión',
            'FUENTE_PRECEDENCIA_D': 'Gate09 C09-01 confirmado Gate10 23-09-2026',
        })
    assert sum(r['CLASIFICACION_INTERNA_C'] != 'POR_REVISAR' for r in rows) == 31
    assert sum(bool(r['VIG_HISTORICO_DESCARTADO_NO_USAR']) for r in rows) == 29
    OUT.mkdir(parents=True, exist_ok=True)
    write_tsv(OUT / 'relacion_estados_situaciones.tsv', rows)
    rule_json = {
        'schema_version': 1, 'version': '2026-10-20.2', 'clasificacion': 'C',
        'proceso': 'Matrícula Unificada', 'edicion': 2026, 'hito': 'H001-RETENCION-IPSS162',
        'grano': 'matrícula/programa + período + corte; no estado global de persona',
        'estado': 'DOCUMENTADA_CONDICIONAL_NO_APLICADA_A_ESTUDIANTES',
        'precedencia': 'D: Gate09 C09-01 confirmado Gate10; no función estado→VIG',
        'campos_de_comprobacion': {
            'nota': 'Conceptos técnicos C; no columnas físicas inventadas de UMAS. SI/NO requiere evidencia y fuente localizable.',
            'contexto': 'SI solo con proceso/edición, identidad de matrícula/programa, período y corte verificados',
            'contradiccion': 'NO solo tras comprobar consistencia temporal de las evidencias',
            'matricula_vigente': 'SI/NO/DESCONOCIDO conforme al reglamento y evidencia del período/corte',
            'condicion_terminal': 'SI: plan finalizado, actividad conducente al título en curso y título aún no obtenido; NO: condición no terminal acreditada; DESCONOCIDO: insuficiente',
        },
        'reglas': [
            {'id': 'R1', 'clasificacion_definicion': 'A', 'clasificacion_operacion': 'C', 'vig': 1,
             'condicion': 'contexto=SI;contradiccion=NO;matricula_vigente=SI;condicion_terminal=NO', 'fuente': 'Manual MU2026 p.31 AF'},
            {'id': 'R2', 'clasificacion_definicion': 'A', 'clasificacion_operacion': 'C', 'vig': 2,
             'condicion': 'contexto=SI;contradiccion=NO;matricula_vigente=SI;condicion_terminal=SI', 'fuente': 'Manual MU2026 p.7 §2.vi;p.31 AF;Gate10 C09-01'},
            {'id': 'R0', 'clasificacion_definicion': 'A', 'clasificacion_operacion': 'C', 'vig': 0,
             'condicion': 'contexto=SI;contradiccion=NO;matricula_vigente=NO', 'fuente': 'Manual MU2026 p.31 AF'},
            {'id': 'RE', 'clasificacion': 'E', 'vig': None, 'condicion': 'evidencia insuficiente o contradictoria; sin default'},
        ],
        'denominadores_preservados_D': {'T2': 222, 'T3': 124},
        'no_redefine': ['retención','continuidad','Avance Curricular','Extranjeros','Titulaciones','elegibilidad de encuestas'],
    }
    write_json(OUT / 'reglas.json', rule_json)
    scenarios = [
        ('Matrícula regular acreditada', 'VIGENTE', 'SI', 'NO', 'SI', 'NO', 'R1', 1),
        ('Proceso terminal completo acreditado', 'EGRESADO', 'SI', 'NO', 'SI', 'SI', 'R2', 2),
        ('Solo etiqueta VIGENTE', 'VIGENTE', 'SI', 'NO', 'DESCONOCIDO', 'DESCONOCIDO', 'RE', None),
        ('Solo etiqueta EGRESADO', 'EGRESADO', 'SI', 'NO', 'DESCONOCIDO', 'DESCONOCIDO', 'RE', None),
        ('Solo etiqueta TITULADO31/42', 'TITULADO', 'SI', 'NO', 'DESCONOCIDO', 'DESCONOCIDO', 'RE', None),
        ('Hoy TITULADO; historial terminal al corte antes de título', 'TITULADO', 'SI', 'NO', 'SI', 'SI', 'R2', 2),
        ('Hoy TITULADO; historial no terminal al corte antes de egreso', 'TITULADO', 'SI', 'NO', 'SI', 'NO', 'R1', 1),
        ('Matrícula confirmada pero terminal desconocido', 'VIGENTE', 'SI', 'NO', 'SI', 'DESCONOCIDO', 'RE', None),
        ('Suspensión y reglamento acreditan no matrícula', 'SUSPENDIDO', 'SI', 'NO', 'NO', 'DESCONOCIDO', 'R0', 0),
        ('Suspensión y reglamento acreditan matrícula no terminal', 'SUSPENDIDO', 'SI', 'NO', 'SI', 'NO', 'R1', 1),
        ('Suspensión y reglamento acreditan matrícula terminal', 'SUSPENDIDO', 'SI', 'NO', 'SI', 'SI', 'R2', 2),
        ('Solo situación16/ELIMINADO', 'ELIMINADO', 'SI', 'NO', 'DESCONOCIDO', 'DESCONOCIDO', 'RE', None),
        ('Contradicción título ya obtenido y actividad para mismo título', 'TITULADO', 'SI', 'SI', 'SI', 'SI', 'RE', None),
        ('Otra matrícula sin resolver grano', 'TITULADO', 'NO', 'NO', 'SI', 'NO', 'RE', None),
        ('Ausencia en extracto parcial no acredita sin matrícula', 'ELIMINADO', 'SI', 'NO', 'DESCONOCIDO', 'DESCONOCIDO', 'RE', None),
    ]
    tested = []
    for label, state, ctx, conflict, matricula, terminal, expected_rule, expected_vig in scenarios:
        evidence = dict(contexto=ctx, contradiccion=conflict, matricula_vigente=matricula, condicion_terminal=terminal)
        rule, vig = decide(evidence)
        assert (rule, vig) == (expected_rule, expected_vig), label
        tested.append({'clasificacion': 'C', 'tipo': 'SINTETICO_NO_ESTUDIANTE', 'escenario': label, 'estado_observado': state,
                       **evidence, 'regla': rule, 'vig': '' if vig is None else vig, 'qa': 'OK'})
    write_tsv(OUT / 'ejemplos_sinteticos.tsv', tested)
    for (_, p, _), entry in zip(sources, inventory):
        assert sha(p) == entry['sha256'], f'Fuente cambió: {p}'
    for path, digest in preserved.items():
        assert sha(ROOT / path) == digest
    qa = {'clasificacion': 'C', 'resultado': 'OK_DOCUMENTAL_CON_PENDIENTES_DE_APLICACION',
          'filas_fuente': 15136, 'pares_observados_y_documentados': 32, 'pares_con_clasificacion_interna': 31,
          'pares_con_definicion_pendiente': 1, 'pares_con_antecedente_historico_descartado': 29,
          'escenarios_sinteticos_ok': len(tested), 'VIG_asignados_a_filas_reales': 0,
          'fuentes_originales_intactas': True, 'base_MU_y_excel_intactos_sha256': preserved,
          'denominadores': {'T2': 222, 'T3': 124}, 'datos_individuales_exportados': False,
          'limite': 'Valida integridad, cobertura de pares y condiciones documentales; no acredita disponibilidad del historial ni carga SIES.'}
    write_json(OUT / 'qa.json', qa)
    write_json(OUT / 'manifest.json', {
        'recurso': 'RELACION_ESTADOS_SITUACIONES_UMAS_VIG_MU2026', 'version': '2026-10-20.2', 'clasificacion': 'C',
        'fecha': datetime.now(ZoneInfo('America/Santiago')).isoformat(), 'repositorio': 'avance_curricular',
        'git_antes': before, 'hito': 'H001-RETENCION-IPSS162',
        'decision_D': 'Usuario solicita crear relación explícita; denominadores T2=222/T3=124 se mantienen. No suministra una nueva equivalencia numérica por estado.',
        'alcance': 'Consulta condicional MU2026; sin asignación a estudiantes ni cambios del motor congelado',
        'fuentes': inventory,
        'salidas_sha256': {p.name: sha(p) for p in sorted(OUT.iterdir()) if p.name != 'manifest.json' and p.is_file()},
        'documentacion': {'ruta': 'docs/relacion_estados_umas_vig_sies_20261020.md', 'sha256': sha(ROOT / 'docs/relacion_estados_umas_vig_sies_20261020.md')},
        'script': {'ruta': str(Path(__file__).relative_to(ROOT)), 'sha256': sha(__file__)},
        'validaciones': qa,
        'pendientes_E': ['Reglamento institucional completo aplicable a vigencia al corte', 'Historial de matrícula y estados fechado UP43/UP66',
                         'Fechas de egreso/título UP59/UP60 y evidencia de actividades terminales', 'Definición/catalogación de situación16/ELIMINADO'],
        'politica': 'No presentar etiquetas ni implementación C como regla A. No VIG por estado aislado. No PII. No commit/push. Snapshot de complementos requiere verificación y nueva publicación si cambia.',
        'siguiente_paso': 'Consumir la relación para explicar condiciones y pendientes; aplicar a matrículas solo con hechos del mismo período/corte acreditados.'
    })
    print(json.dumps(qa, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
