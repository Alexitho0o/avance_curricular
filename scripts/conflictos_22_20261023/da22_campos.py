from pathlib import Path
from openpyxl import load_workbook
import json,sqlite3,collections,hashlib,itertools
P=Path('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023')
c=sqlite3.connect('file:/Users/alexi/.local/share/ipss/cierre_inf35_03_t3_20261023_154224/DICCIONARIO_SIES.sqlite?mode=ro',uri=True)
rs=[json.loads(x[0]) for x in c.execute('select payload from observaciones')];annual={}
for r in rs:
 if r['codigo_documental_atribuible']:
  k=tuple(str(r[f]) for f in ['persona','CODCLI','plan','carrera','sede_dimension','jornada_sies','ano'])
  if k not in annual:annual[k]=r
G=collections.defaultdict(list)
for r in annual.values():G[tuple(str(r[f]) for f in ['carrera','sede_dimension','jornada_sies','ano','plan'])].append(r)
G={k:v for k,v in G.items() if len({r['codigo_documental_atribuible'] for r in v})>1};assert len(G)==22
refs=collections.defaultdict(lambda:collections.defaultdict(set))
for gr in G.values():
 for r in gr:
  for d in r['contextos_DA']:refs[d['fuente']][d['hoja']].add(int(d['fila']))
cache={};sources=[]
for orig,shs in refs.items():
 p=Path(orig)
 if not p.exists():p=next(Path('/Users/alexi/Library/CloudStorage').rglob(p.name))
 sources.append({'ruta_original':orig,'ruta_actual':str(p),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()});w=load_workbook(p,read_only=True,data_only=True)
 for sh,needed in shs.items():
  s=w[sh];hh=next(s.iter_rows(min_row=1,max_row=1,values_only=True))
  for row,vals in enumerate(s.iter_rows(min_row=2,max_row=max(needed),values_only=True),2):
   if row in needed:cache[(orig,sh,row)]=dict(zip(hh,vals))
 w.close()
fields=['CATEGORIA','VIASDEADMISION','SITUACION','ESTADOACADEMICO'];out=[]
for g,gr in sorted(G.items()):
 rr=[]
 for r in gr:
  ds=[cache[(d['fuente'],d['hoja'],int(d['fila']))] for d in r['contextos_DA']]
  assert all(str(d['CODCLI']).strip()==str(r['CODCLI']).strip() for d in ds)
  rr.append({'persona':r['persona'],'codigo':r['codigo_documental_atribuible'],'vals':{f:sorted({str(d.get(f) or '').strip() for d in ds}) for f in fields}})
 tests=[]
 for fset in [[f] for f in fields]+[fields[:2],fields]:
  mp=collections.defaultdict(set);count=collections.defaultdict(lambda:collections.defaultdict(set));missing=0;multi=0
  for r in rr:
   if any(not r['vals'][f] or r['vals'][f]==[''] for f in fset):missing+=1;continue
   if any(len(r['vals'][f])!=1 for f in fset):multi+=1;continue
   v=tuple(r['vals'][f][0] for f in fset);mp[v].add(r['codigo']);count[r['codigo']][' | '.join(v)].add(r['persona'])
  tests.append({'variables':fset,'separacion_completa_descriptiva':missing==multi==0 and all(len(x)==1 for x in mp.values()),'contextos_sin_valor':missing,'contextos_valor_multiple':multi,'valores_compartidos_entre_codigos':sum(len(x)>1 for x in mp.values()),'distribucion_personas_por_codigo_valor':{c:{v:len(p) for v,p in m.items()} for c,m in count.items()}})
 out.append({'grupo':dict(zip(['carrera','sede','jornada','anio','plan'],g)),'contextos':len(gr),'personas':len({r['persona'] for r in gr}),'pruebas':tests})
result={'clasificacion':'B campos actuales observados; C separación descriptiva; E no acredita historia ni regla de versión','grupos':out,'fuentes':sources,'QA':{'22_grupos':len(out)==22,'filas_DA_con_CODCLI_correcto':True,'contextos':sum(x['contextos'] for x in out),'estado_corte_no_historico':True,'base_modificada':False},'limite':'Estados y situación provienen del corte DatosAlumnos disponible, no del período histórico analizado; pueden ser consecuencia posterior. Separación descriptiva no es validación fuera de persona ni adopción.'}
text=json.dumps(result,ensure_ascii=False,indent=2)+'\n'
for gr in G.values():
 for r in gr:assert r['persona'] not in text and r['CODCLI'] not in text
Path('/tmp/da22_campos_agregado.json').write_text(text)
print(json.dumps({'file':'/tmp/da22_campos_agregado.json','fuentes':sources,'separaciones':[{'grupo':x['grupo'],'variables':[p['variables'] for p in x['pruebas'] if p['separacion_completa_descriptiva']]} for x in out if any(p['separacion_completa_descriptiva'] for p in x['pruebas'])]},ensure_ascii=False,indent=2))
