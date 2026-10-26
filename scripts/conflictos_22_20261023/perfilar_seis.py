from pathlib import Path
from openpyxl import load_workbook
import json,collections,hashlib,sqlite3
P=Path('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023');E=json.loads((P/'ERRORES6_DETALLE.json').read_text())
muallow=['ANIO_INFORMADO','COD_SED','COD_CAR','MODALIDAD','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','ASI_INS_ANT','ASI_APR_ANT','ASI_INS_HIS','ASI_APR_HIS','NIV_ACA','FECHA_MATRICULA','REINCORPORACION','VIG','MM_ANIO_ORIGEN','MM_SEM_ORIGEN','MM_ANIO_ACTUAL','MM_SEM_ACTUAL','MM_SEM_SUSPENSION','MM_EST_PROC_TER_TIT','MM_ANIO_EGRESO','MM_SEM_EGRESO','MC_CODIGO_UNICO','MC_MODALIDAD','MC_JORNADA','MC_TIPO_PLAN_CARR','MC_CARAC_PLAN_ESP','MC_DUR_ESTUDIO_CARR','MC_DUR_PROCESO_TIT','MC_DUR_TOTAL_CARR','MC_NIVEL_GLOBAL','MC_NIVEL_CARRERA_1']
daallow=['ANOMATRICULA','PERIODOMATRICULA','ANOINGRESO','PERIODOINGRESO','FECHAMATRICULA','JORNADA','CATEGORIA','VIASDEADMISION','ESTADOACADEMICO','SITUACION','NIVEL','SEDE']
refs=collections.defaultdict(lambda:collections.defaultdict(set))
for r in E:
 for t in r['all_group']:
  for o in t['obs']:
   for f in ['contextos_DA','evidencias_matricula']:
    for d in o[f]:refs[d['fuente']][d['hoja']].add(int(d['fila']))
cache={};source=[]
for original,shs in refs.items():
 p=Path(original)
 if not p.exists():p=next(Path('/Users/alexi/Library/CloudStorage').rglob(p.name))
 h=hashlib.sha256(p.read_bytes()).hexdigest();source.append({'fuente_original':original,'ruta_actual':str(p),'sha256':h})
 w=load_workbook(p,read_only=True,data_only=True)
 for sh,rows in shs.items():
  s=w[sh];head=next(s.iter_rows(min_row=1,max_row=1,values_only=True))
  for rown,vals in enumerate(s.iter_rows(min_row=2,max_row=max(rows),values_only=True),2):
   if rown in rows:cache[(original,sh,rown)]=dict(zip(head,vals))
 w.close()
# QA de localizadores, sin imprimir identificadores. Cabeceras históricas explícitas.
def identval(v):
 return str(v if v is not None else '').strip().removesuffix('.0').upper()
qa_identity={'referencias_DA_verificadas_CODCLI':0,'referencias_MU_verificadas_documento_DV':0,'referencias_MU_verificadas_TIPO_DOC':0,'referencias_MU_sin_TIPO_DOC_fuente':0,'referencias_MU_sin_TIPO_DOC_payload':0,'fallos':0,'mapeos':{'MU_derivada':['TIPO_DOC','N_DOC','DV'],'MU_historica':['TIPO_DOC_si_existe','MM_N_DOCUMENTO','MM_DV'],'DA':['CODCLI']}}
seen=set()
for r in E:
 for t in r['all_group']:
  for o in t['obs']:
   for d in o['contextos_DA']:
    ref=(d['fuente'],d['hoja'],int(d['fila']))
    if ('DA',*ref) in seen:continue
    seen.add(('DA',*ref));raw=cache[ref]
    assert identval(raw['CODCLI'])==identval(d['CODCLI'])==identval(t['CODCLI'])
    qa_identity['referencias_DA_verificadas_CODCLI']+=1
   for d in o['evidencias_matricula']:
    ref=(d['fuente'],d['hoja'],int(d['fila']))
    if ('MU',*ref) in seen:continue
    seen.add(('MU',*ref));raw=cache[ref]
    docfield='N_DOC' if 'N_DOC' in raw else 'MM_N_DOCUMENTO'
    dvfield='DV' if 'DV' in raw else 'MM_DV'
    assert identval(raw[docfield])==identval(d['documento_original'])
    assert identval(raw[dvfield])==identval(d['dv_original'])
    qa_identity['referencias_MU_verificadas_documento_DV']+=1
    if 'TIPO_DOC' not in raw:
     qa_identity['referencias_MU_sin_TIPO_DOC_fuente']+=1
    elif not identval(d.get('tipo_doc_original')):
     qa_identity['referencias_MU_sin_TIPO_DOC_payload']+=1
    else:
     assert identval(raw['TIPO_DOC'])==identval(d['tipo_doc_original'])
     qa_identity['referencias_MU_verificadas_TIPO_DOC']+=1
Path('/tmp/errores6_identidad_filas_QA.json').write_text(json.dumps(qa_identity,ensure_ascii=False,indent=2))
for r in E:
 for t in r['all_group']:
  o=t['obs'][0]
  t['da_raw']=[{'ref':{'ruta':d['fuente'],'hoja':d['hoja'],'fila':d['fila']},'campos':{k:str(v or '') for k,v in cache[(d['fuente'],d['hoja'],int(d['fila']))].items() if k in daallow}} for d in o['contextos_DA']]
  t['mu_raw']=[{'ref':{'ruta':d['fuente'],'hoja':d['hoja'],'fila':d['fila']},'campos':{k:str(v if v is not None else '') for k,v in cache[(d['fuente'],d['hoja'],int(d['fila']))].items() if k in muallow}} for d in o['evidencias_matricula']]
  states=collections.Counter(f['estado_ramo'] for oo in t['obs'] for f in oo['fuentes_promedios']);t['actividad_ramos']=dict(states)
(P/'ERRORES6_PERFILES_PRIVADOS.json').write_text(json.dumps(E,ensure_ascii=False,indent=2,default=str))
Path('/tmp/errores6_fuentes.json').write_text(json.dumps(source,ensure_ascii=False,indent=2))
out=[]
for r in E:
 target=next(t for t in r['all_group'] if t['persona']==r['persona']);others=[t for t in r['all_group'] if t['persona']!=r['persona']]
 item={'id':r['id'],'codigo':r['codigo'],'pred':r['pred'],'objetivo_DA':target['da_raw'],'objetivo_MU':target['mu_raw'],'actividad_objetivo':target['actividad_ramos'],'contrastes':[]}
 for typ,fields in [('da_raw',daallow),('mu_raw',muallow)]:
  for k in fields:
   tv={d['campos'].get(k,'') for d in target[typ]};oc=collections.Counter(d['campos'].get(k,'') for t in others for d in t[typ])
   if not tv:continue
   if set(oc)!=tv:item['contrastes'].append({'fuente':typ,'variable':k,'objetivo':sorted(tv),'otros':dict(oc),'separacion_completa':not (tv&set(oc))})
 out.append(item)
Path('/tmp/errores6_perfiles_agregados.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str))
print(json.dumps(out,ensure_ascii=False,indent=2,default=str))
