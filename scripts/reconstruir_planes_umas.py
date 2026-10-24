#!/usr/bin/env python3
"""Audita estructura y asociaciones UMAS sin inferir vigencia ni VERSION SIES.

Contrato: docs/fase_4_planes_asignaciones_umas.md. XLSX solo lectura; identidades
solo en memoria. Salidas agregadas con localizadores, sin notas ni identidades.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import sys

PROM_FIELDS=('CODCLI','CODCARR','CARRERA','JORNADA','REGIMEN','CODRAMO','RAMOEQUIV',
             'ANO','PERIODO','NIVEL','PLAN_DE_ESTUDIO')
ALUM_FIELDS=('CODCLI','CODCARPR','NOMBRE_L','JORNADA','SEDE','ANOMATRICULA',
             'PERIODOMATRICULA','ANOINGRESO','PERIODOINGRESO')
OBS_FIELDS=('PLAN_DE_ESTUDIO','CODCARR','CARRERA','JORNADA','ANO','PERIODO','REGIMEN')
MISSING={'','NULL','NONE','NAN'}


def digest(path):
    with Path(path).open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()


def value(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v)


def token(v):
    """Solo espacios periféricos. NULL/vacíos nunca se unen como claves."""
    s=value(v).strip()
    return '' if s.upper() in MISSING else s


def json_bytes(v):
    return (json.dumps(v,ensure_ascii=False,indent=2,sort_keys=True)+'\n').encode()


def tsv_bytes(rows):
    fields=list(dict.fromkeys(k for r in rows for k in r))
    out=io.StringIO(newline='');w=csv.DictWriter(out,fieldnames=fields,delimiter='\t',lineterminator='\n')
    w.writeheader()
    for r in rows:
        w.writerow({k:json.dumps(v,ensure_ascii=False,sort_keys=True) if isinstance(v,(list,dict)) else v for k,v in r.items()})
    return out.getvalue().encode()


def read_sheet(path,sheet,fields=None):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        it=wb[sheet].iter_rows(values_only=True);headers=[value(v) for v in next(it)]
        selected=list(fields) if fields is not None else headers
        for f in selected:
            if headers.count(f)!=1:raise ValueError('Campo ausente o duplicado: '+f)
        positions={f:headers.index(f) for f in selected};rows=[]
        for i,values in enumerate(it,2):
            if any(v is not None for v in values):
                rows.append({**{f:(values[j] if fields is None else value(values[j]))
                                for f,j in positions.items()},'_fila':i})
        return rows,{'sheet':sheet,'headers':headers,'filas_datos':len(rows),
                     'columnas_proyectadas':selected,'fila_encabezado':1,
                     'valores':'Tipos originales del catálogo' if fields is None else
                               'Proyección académica textual; None a vacío, enteros sin .0; claves eliminan padding y excluyen NULL.'}
    finally:wb.close()


def observation_state(r,start=2007,end=2026):
    y=token(r['ANO'])
    if not re.fullmatch(r'\d{4}',y):return 'ANIO_NO_VALIDO'
    if not start<=int(y)<=end:return 'FUERA_PERIODO'
    if any(not token(r[f]) for f in ('CODCLI','CODCARR','PLAN_DE_ESTUDIO','CODRAMO','PERIODO')):
        return 'CLAVES_ACADEMICAS_INCOMPLETAS'
    return 'OBSERVACION_CON_ANIO'


def equality(a,b):
    a,b=token(a),token(b)
    return 'SIN_DATO' if not a or not b else 'IGUAL' if a==b else 'DIFERENTE'


def year_tokens(plan):
    """Solo patrón textual, deliberadamente sin atribuirle semántica temporal."""
    return sorted(set(re.findall(r'(?:19|20)\d{2}',plan)))


def analyze(plan_rows,prom,alum,start=2007,end=2026):
    """Enlace de contexto de reportes; nunca asignación efectiva en un año."""
    plan_contexts=defaultdict(set);plan_courses=defaultdict(set);plan_equiv=defaultdict(set)
    plan_names=defaultdict(set);plan_first={}
    for r in plan_rows:
        p=token(r['CODPESTUD'])
        if not p:raise ValueError('Plan vacío en catálogo')
        plan_contexts[p].add((token(r['Codcarr']),token(r['Sede'])))
        plan_names[p].add(r['Nombre_L']);plan_courses[p].add(token(r['CODRAMO']))
        e=token(r['RamoEquiv'])
        # El listado declara PLAN_DESTINO/RAMO_DESTINO. La nota no acredita
        # el plan de origen de ese ramo: solo se registra coincidencia en relación.
        if e and e.count('/')==1 and ':' not in e:
            destination_plan,destination_course=e.split('/')
            if token(destination_plan) and token(destination_course):
                plan_equiv[p].add(token(destination_course))
        plan_first[p]=min(plan_first.get(p,r['_fila']),r['_fila'])
    da=defaultdict(list)
    for r in alum:
        if token(r['CODCLI']):da[token(r['CODCLI'])].append(r)
    person_plans=defaultdict(set);person_codes=defaultdict(set);person_years=defaultdict(set)
    for r in prom:
        pid=token(r['CODCLI']);p=token(r['PLAN_DE_ESTUDIO'])
        if pid:
            if p:person_plans[pid].add(p)
            if token(r['CODCARR']):person_codes[pid].add(token(r['CODCARR']))
            if observation_state(r,start,end)=='OBSERVACION_CON_ANIO':person_years[pid].add(token(r['ANO']))
    grouped={};join_groups={};temporal_groups={};plan_obs=defaultdict(list)
    statuses=Counter();joins=Counter();course_statuses=Counter();career_statuses=Counter();sede_statuses=Counter();schedule_statuses=Counter()
    da_career_statuses=Counter();da_year_statuses=Counter()
    def add_group(bucket,key,base,row,pid):
        if key not in bucket:bucket[key]={**base,'filas_reporte':0,'_personas':set(),'fila_promedios_ejemplo':row['_fila']}
        g=bucket[key];g['filas_reporte']+=1
        if pid:g['_personas'].add(pid)
        g['fila_promedios_ejemplo']=min(g['fila_promedios_ejemplo'],row['_fila'])
        return g
    for r in prom:
        p=token(r['PLAN_DE_ESTUDIO']);pid=token(r['CODCLI']);state=observation_state(r,start,end);statuses[state]+=1
        base={f:r[f] for f in OBS_FIELDS};base['estado_observacion']=state
        add_group(grouped,tuple(base.values()),base,r,pid)
        if p:plan_obs[p].append(r)
        contexts=plan_contexts.get(p,set());pc=next(iter(contexts)) if len(contexts)==1 else None
        ps='PLAN_NO_LOCALIZADO' if not contexts else 'PLAN_CONTEXTO_AMBIGUO' if len(contexts)>1 else 'PLAN_CONTEXTO_UNICO'
        match_plan_career=equality(r['CODCARR'],pc[0]) if pc else 'NO_COMPARABLE'
        career_statuses[match_plan_career]+=1
        course=token(r['CODRAMO'])
        cs=('RAMO_AUSENTE' if not course else 'PLAN_NO_LOCALIZADO' if not contexts else
            'RAMO_EN_PLAN' if course in plan_courses[p] else 'RAMO_APARECE_EN_EQUIVALENCIA_CATALOGADA' if course in plan_equiv[p] else
            'RAMO_NO_LOCALIZADO_EN_PLAN')
        course_statuses[cs]+=1
        candidates=da.get(pid,[]) if pid else []
        ds='SIN_IDENTIFICADOR' if not pid else 'SIN_COINCIDENCIA' if not candidates else 'MULTIPLES_FILAS' if len(candidates)>1 else 'CONTEXTO_UNICO_DEL_REPORTE'
        d=candidates[0] if len(candidates)==1 else None
        dy=token(d['ANOMATRICULA']) if d else ''
        ds_year='NO_COMPARABLE' if not d else 'ANIO_NO_VALIDO' if not re.fullmatch(r'\d{4}',dy) else 'DENTRO_AMBITO' if start<=int(dy)<=end else 'FUERA_AMBITO'
        dc=equality(r['CODCARR'],d['CODCARPR']) if d else 'NO_COMPARABLE'
        dj=equality(r['JORNADA'],d['JORNADA']) if d else 'NO_COMPARABLE'
        sede=equality(pc[1],d['SEDE']) if d and pc else 'NO_COMPARABLE'
        joins[ds]+=1;schedule_statuses[dj]+=1;sede_statuses[sede]+=1
        da_career_statuses[dc]+=1;da_year_statuses[ds_year]+=1
        # Se compara contexto del corte incluso si es 2027, pero se marca fuera
        # del ámbito y nunca se adjudica a la observación histórica.
        join={**base,'estado_plan_catalogo':ps,'carrera_plan_vs_promedios':match_plan_career,
              'estado_ramo':cs,'estado_datos_alumnos':ds,'ambito_anio_matricula_da':ds_year,
              'carrera_da_vs_promedios':dc,'jornada_da_vs_promedios':dj,'sede_da_vs_plan':sede,
              'CODCARPR_DA':d['CODCARPR'] if d else '', 'JORNADA_DA':d['JORNADA'] if d else '',
              'SEDE_DA':d['SEDE'] if d else '', 'ANOMATRICULA_DA':d['ANOMATRICULA'] if d else '',
              'ANOINGRESO_DA':d['ANOINGRESO'] if d else '',
              'CODCARR_LISTADO':pc[0] if pc else '', 'SEDE_LISTADO':pc[1] if pc else ''}
        g=add_group(join_groups,tuple(join.values()),join,r,pid)
        if g['fila_promedios_ejemplo']==r['_fila']:
            g['fila_alumnos_ejemplo']=d['_fila'] if d else ''
        g.setdefault('fila_plan_ejemplo',plan_first.get(p,''))
        if state=='OBSERVACION_CON_ANIO':
            ts=year_tokens(p)
            relation='SIN_PATRON_ANUAL' if not ts else 'VARIOS_PATRONES_ANUALES' if len(ts)>1 else (
                'ACTIVIDAD_ANTERIOR_AL_PATRON' if int(r['ANO'])<int(ts[0]) else
                'ACTIVIDAD_MISMO_ANIO_DEL_PATRON' if int(r['ANO'])==int(ts[0]) else 'ACTIVIDAD_POSTERIOR_AL_PATRON')
            t={'PLAN_DE_ESTUDIO':r['PLAN_DE_ESTUDIO'],'CODCARR':r['CODCARR'],'ANO':r['ANO'],
               'PERIODO':r['PERIODO'],'patrones_anuales_textuales':','.join(ts),'relacion_textual':relation,
               'significado_temporal':'NO_ACREDITADO'}
            add_group(temporal_groups,tuple(t.values()),t,r,pid)
    def public_groups(groups):
        out=[]
        for _,r in sorted(groups.items()):
            out.append({**{k:v for k,v in r.items() if not k.startswith('_')},'identificadores_internos_distintos':len(r['_personas'])})
        return out
    coverage=[]
    for p in sorted(set(plan_contexts)|set(plan_obs)):
        rs=plan_obs.get(p,[]);ys=sorted({token(r['ANO']) for r in rs if observation_state(r,start,end)=='OBSERVACION_CON_ANIO'})
        coverage.append({'CODPESTUD':p,'presencia':'CATALOGO_Y_PROMEDIOS' if p in plan_contexts and rs else 'SOLO_CATALOGO' if p in plan_contexts else 'SOLO_PROMEDIOS',
                         'contextos_catalogo':[{'Codcarr':c,'Sede':s} for c,s in sorted(plan_contexts.get(p,set()))],
                         'nombres_internos_catalogo':sorted(plan_names.get(p,set())),
                         'carreras_promedios':sorted({r['CODCARR'] for r in rs}),
                         'jornadas_promedios':sorted({r['JORNADA'] for r in rs}),
                         'anios_actividad_observada':ys,'filas_promedios':len(rs),
                         'filas_con_anio_y_claves':sum(observation_state(r,start,end)=='OBSERVACION_CON_ANIO' for r in rs),
                         'identificadores_internos_distintos':len({token(r['CODCLI']) for r in rs if token(r['CODCLI'])}),
                         'fila_plan_ejemplo':plan_first.get(p,''),'fila_promedios_ejemplo':min((r['_fila'] for r in rs),default=''),
                         'vigencia_desde_acreditada':None,'vigencia_hasta_acreditada':None,
                         'asignacion_temporal_acreditada':False})
    temporal=public_groups(temporal_groups);joined=public_groups(join_groups)
    repeats=[]
    for rows in da.values():
        if len(rows)>1:
            repeats.append({'filas_alumnos':sorted(r['_fila'] for r in rows),
                            'contextos_academicos':[ {f:r[f] for f in ALUM_FIELDS if f!='CODCLI'} for r in rows],
                            'estado':'NO_SELECCIONAR_UNA_FILA'})
    valid=[r for r in prom if observation_state(r,start,end)=='OBSERVACION_CON_ANIO']
    summary={'fase':4,'estado':'RECONSTRUIDO_CON_BRECHAS_TEMPORALES','periodo_objetivo':[start,end],
             'filas_promedios':len(prom),'filas_datos_alumnos':len(alum),'filas_listado_planes':len(plan_rows),
             'planes_catalogo':len(plan_contexts),'planes_promedios':len(plan_obs),'planes_union':len(coverage),
             'planes_promedios_con_anio_y_claves':sum(r['filas_con_anio_y_claves']>0 for r in coverage),
             'planes_promedios_solo_observaciones_sin_anio_o_claves':sum(r['filas_promedios']>0 and r['filas_con_anio_y_claves']==0 for r in coverage),
             'planes_catalogo_sin_actividad_con_anio_y_claves':sum(r['presencia']!='SOLO_PROMEDIOS' and r['filas_con_anio_y_claves']==0 for r in coverage),
             'presencia_planes':dict(Counter(r['presencia'] for r in coverage)),
             'planes_contexto_catalogo_multiple':sum(len(v)>1 for v in plan_contexts.values()),
             'estados_observacion_filas':dict(statuses),'anios_actividad_con_claves':sorted({int(r['ANO']) for r in valid}),
             'anios_objetivo_sin_actividad_observada':sorted(set(range(start,end+1))-{int(r['ANO']) for r in valid}),
             'identificadores_promedios_distintos':len({token(r['CODCLI']) for r in prom if token(r['CODCLI'])}),
             'identificadores_con_varios_planes_en_reporte':sum(len(v)>1 for v in person_plans.values()),
             'identificadores_con_varias_carreras_en_reporte':sum(len(v)>1 for v in person_codes.values()),
             'identificadores_con_actividad_varios_anios':sum(len(v)>1 for v in person_years.values()),
             'identificadores_datos_alumnos_distintos':len(da),'identificadores_datos_alumnos_repetidos':len(repeats),
             'filas_datos_alumnos_por_anio_matricula':dict(sorted(Counter(r['ANOMATRICULA'] for r in alum).items())),
             'identificadores_compartidos_promedios_da':len({token(r['CODCLI']) for r in prom if token(r['CODCLI'])}&set(da)),
             'estados_enlace_da_por_fila_promedios':dict(joins),'carrera_plan_vs_promedios_por_fila':dict(career_statuses),
             'carrera_da_vs_promedios_por_fila':dict(da_career_statuses),
             'ambito_matricula_da_por_fila_promedios':dict(da_year_statuses),
             'jornada_da_vs_promedios_por_fila':dict(schedule_statuses),'sede_da_vs_plan_por_fila':dict(sede_statuses),
             'ramo_vs_plan_por_fila':dict(course_statuses),
             'actividad_anterior_al_patron_anual_en_codigo_plan':{
                 'filas':sum(g['filas_reporte'] for g in temporal if g['relacion_textual']=='ACTIVIDAD_ANTERIOR_AL_PATRON'),
                 'planes':len({g['PLAN_DE_ESTUDIO'] for g in temporal if g['relacion_textual']=='ACTIVIDAD_ANTERIOR_AL_PATRON'}),
                 'interpretacion':'Patrón lexical, no fecha de creación ni prueba de migración.'},
             'asignaciones_temporales_certificadas':0,
             'limites':['La asociación CODCLI/plan en el reporte no demuestra asignación del plan en cada año de la nota.',
                       'Las uniones comparan contextos de reportes con corte real no acreditado.',
                       'No se deduce modalidad de jornada ni de MODALIDAD escolar de DatosAlumnos.',
                       'Ausencia de ramo requiere revisión; no prueba error, equivalencia ni convalidación.',
                       'Coincidencia de ramo en relación PLAN_DESTINO/RAMO_DESTINO no acredita homologación individual ni plan de origen de la nota.',
                       'Los identificadores internos distintos no se presentan como personas únicas institucionales.',
                       'No hay enlace individual MRUN–UMAS ni equivalencia plan–VERSION SIES certificada.']}
    examples=[]
    for c in ('ILOG','CILOG','NATP','CNATP','IECIRE','IECSRE','ASREOL'):
        examples.append({'codigo_carrera_interno':c,'planes':[r for r in coverage if any(x['Codcarr']==c for x in r['contextos_catalogo'])],
                         'observaciones_temporales':[r for r in temporal if r['CODCARR']==c],
                         'nota':'Ejemplo de programa; no identidad ni versión SIES asignada.'})
    return {'cobertura_planes.tsv':coverage,'observaciones_plan_periodo.tsv':public_groups(grouped),
            'conciliacion_contextos_umas.tsv':joined,'diagnostico_patrones_anuales.tsv':temporal,
            'duplicados_datos_alumnos.json':repeats,'ejemplos_academicos.json':examples,'resumen_asociaciones.json':summary}


def run(config_path,output):
    config=json.loads(config_path.read_text());sources=config['sources'];paths={k:Path(s['path']) for k,s in sources.items()}
    if config['periodo_objetivo']!=[2007,2026]:raise ValueError('Contrato H4: ámbito explícito 2007–2026')
    for k,p in paths.items():
        if digest(p)!=sources[k]['sha256']:raise ValueError('Fuente alterada: '+k)
    pr,ps=read_sheet(paths['planes'],sources['planes']['sheet'])
    nr,ns=read_sheet(paths['promedios'],sources['promedios']['sheet'],PROM_FIELDS)
    ar,ass=read_sheet(paths['alumnos'],sources['alumnos']['sheet'],ALUM_FIELDS)
    outputs=analyze(pr,nr,ar)
    # Auxiliar puro sobre estructura, sin datos de estudiantes.
    sys.path.insert(0,str(Path(__file__).parent))
    from estructura_planes_umas import build_catalog
    catalog=build_catalog(pr)
    outputs.update({'catalogo_planes.json':catalog['catalogo_planes'],
                    'estructura_asignaturas.json':catalog['estructura_asignaturas'],
                    'comparaciones_planes.json':catalog['comparaciones_planes'],
                    'resumen_catalogo.json':catalog['resumen']})
    outputs['esquemas_lectura.json']={'planes':ps,'promedios':ns,'alumnos':ass}
    summary={'fase':4,'estado':'COMPLETADA_CON_BRECHAS_TEMPORALES','asociaciones':outputs['resumen_asociaciones.json'],
             'catalogo':catalog['resumen'],'fuentes_documentales':'evidencia_temporal.json',
             'paso_siguiente':'Fase5: reglas con evidencia suficiente; asignación temporal y casos sin catálogo permanecen pendientes.'}
    outputs['resumen.json']=summary
    checks={'filas_promedios_conciliadas':sum(r['filas_reporte'] for r in outputs['observaciones_plan_periodo.tsv'])==len(nr),
            'enlace_sin_expansion':sum(r['filas_reporte'] for r in outputs['conciliacion_contextos_umas.tsv'])==len(nr),
            'planes_preservados':len(outputs['cobertura_planes.tsv'])==len({token(r['CODPESTUD']) for r in pr}|{token(r['PLAN_DE_ESTUDIO']) for r in nr if token(r['PLAN_DE_ESTUDIO'])}),
            'fuentes_sin_modificar':all(digest(p)==sources[k]['sha256'] for k,p in paths.items())}
    if not all(checks.values()):raise ValueError('Fallo de conciliación')
    data={n:tsv_bytes(v) if n.endswith('.tsv') else json_bytes(v) for n,v in outputs.items()}
    data['fuentes.json']=config_path.read_bytes()
    data['manifest.json']=json_bytes({'fase':4,'scripts_sha256':{p.name:digest(p) for p in [Path(__file__),Path(__file__).with_name('estructura_planes_umas.py')]},
                                    'checks':checks,'outputs':{n:hashlib.sha256(v).hexdigest() for n,v in sorted(data.items())},
                                    'privacidad':'Sin identidades o notas individuales; enlaces internos en memoria y salidas académicas agregadas.'})
    for n,b in data.items():
        if (output/n).exists() and (output/n).read_bytes()!=b:raise ValueError('Salida existente distinta: '+str(output/n))
    output.mkdir(parents=True,exist_ok=True)
    for n,b in data.items():
        if not (output/n).exists():(output/n).write_bytes(b)
    return summary


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--fuentes',required=True,type=Path);p.add_argument('--salida',required=True,type=Path)
    a=p.parse_args();print(json.dumps(run(a.fuentes,a.salida),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
