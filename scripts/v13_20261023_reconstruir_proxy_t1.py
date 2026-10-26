"""C: reproduce el proxy heredado sin adjudicar versiones ni modificar fuentes.

E: el estado del extracto PROMEDIOS no acredita vigencia histórica en un cierre.
La función devuelve datos individuales solo en memoria; la CLI exporta únicamente
controles agregados. Las fuentes privadas y sus hashes siguen fijados al corte.
"""
from pathlib import Path
from collections import defaultdict, Counter
import argparse
import hashlib
import json
import openpyxl


def reconstruir_proxy(repo_root: Path):
    R=repo_root.resolve();M=json.loads((R/'outputs/aplicacion_diccionario_intertrimestral_20261023_125556/MANIFEST_APLICACION.json').read_text());src=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx');prom=Path('/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_12166.xlsx');sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    assert sha(src)==M['fuentes'][str(src)];assert sha(prom)=='aad1a71d166bdfb839049a5db534a21d050df1ed4e9251ad9982245896429015'
    op=Path(M['privado'])/'ORIGENES_T1.json';assert sha(op)==M['auditoria_privada_hashes']['ORIGENES_T1.json'];origins=json.loads(op.read_text())
    w=openpyxl.load_workbook(src,read_only=True,data_only=True);byrow={i:r for i,r in enumerate(w['20_Foto_BBDD'].iter_rows(values_only=True),1)}
    keys={};people=defaultdict(set)
    for o in origins:
     r=byrow[o['fila_foto']];assert r[1]==o['anio'] and r[2]==1;assert str(r[9])==o['carrera'];key=str(int(r[8]));assert (o['anio'],key) not in keys;keys[(o['anio'],key)]=o;people[o['anio']].add(key)
    anyrows=defaultdict(set);vig=defaultdict(set);rawrows=Counter();invalid=Counter()
    def body(v):
     if v is None:return None
     if isinstance(v,(int,float)):return str(int(v))
     t=str(v).strip().replace('.','').split('-')[0]
     if t.isdigit():return str(int(t))
     return None
    p=openpyxl.load_workbook(prom,read_only=True,data_only=True);s=p['Hoja1'];h=list(next(s.iter_rows(max_row=1,values_only=True)));assert all(k in h for k in ['RUT','DIG','CODCLI','ANO','PERIODO','ESTADO_ACADEMICO'])
    for rn,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
     d=dict(zip(h,row));y=d['ANO'];t=d['PERIODO'];k=body(d['RUT']);rawrows[(y,t)]+=1
     if k is None:invalid[(y,t)]+=1;continue
     anyrows[(y,t)].add(k)
     if str(d['ESTADO_ACADEMICO']).strip()=='VIGENTE':vig[(y,t)].add(k)
    p.close()
    out=[]
    for r in [1850,1851,1852,1853,1854,1855]:
     rr=next(w['26_Avance_Curricular_Resumen'].iter_rows(min_row=r,max_row=r,values_only=True));cohort=int(rr[0]);y,t=map(int,rr[1].split('-'));obs=len(people[cohort]&anyrows[(y,t)]);n=len(people[cohort]&vig[(y,t)]);out.append({'fila26':r,'cohorte':cohort,'anio_actividad':y,'periodo':t,'denom':len(people[cohort]),'denom_heredado':rr[2],'observados':obs,'observados_heredado':rr[3],'vigentes':n,'vigentes_heredado':rr[4],'coincide':obs==rr[3] and n==rr[4] and len(people[cohort])==rr[2]})
    w.close()
    return {name: value for name, value in locals().items() if name in {"R", "M", "src", "sha", "origins", "byrow", "vig", "anyrows", "keys", "op", "prom", "out"}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output-dir", type=Path, required=True,
                        help="Carpeta para controles agregados; no escribe detalles personales.")
    args = parser.parse_args()
    target = args.output_dir / "proxy_T1_reconstruccion_controles.json"
    if target.exists():
        parser.error(f"La salida ya existe: {target}")
    result = reconstruir_proxy(args.repo_root)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result["out"], ensure_ascii=False, indent=2))
    print(json.dumps(result["out"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
