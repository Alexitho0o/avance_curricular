#!/usr/bin/env python3

from pathlib import Path
from datetime import datetime
import hashlib
import json
import re
import unicodedata

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RAIZ = Path(
    "/Users/alexi/Documents/GitHub/avance_curricular"
).resolve()

EJECUCION = (
    RAIZ
    / "avance_curricular_2026"
    / "05_cierre_integral"
    / "CIERRE_AVANCE_CURRICULAR_20260703_125157"
)

EXPEDIENTE = (
    EJECUCION
    / "03_CONCILIACION_CARRERAS"
    / "EXPEDIENTE_42_PERIODIZACIONES_20260703_134419"
)

FUENTE_EXCEL = (
    EXPEDIENTE
    / "02_RESULTADOS"
    / "EXPEDIENTE_42_PERIODIZACIONES.xlsx"
)

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

SALIDA = (
    EXPEDIENTE
    / f"SOLICITUD_PERIODIZACION_15_PLANES_{timestamp}"
)

RESULTADOS = SALIDA / "02_RESULTADOS"
AUDITORIA = SALIDA / "03_AUDITORIA"
ENTREGA = SALIDA / "04_ENTREGA_INSTITUCIONAL"

for carpeta in (
    RESULTADOS,
    AUDITORIA,
    ENTREGA,
):
    carpeta.mkdir(
        parents=True,
        exist_ok=True,
    )


def norm_col(valor):
    texto = str(valor or "").strip().upper()

    texto = unicodedata.normalize(
        "NFKD",
        texto,
    )

    texto = "".join(
        caracter
        for caracter in texto
        if not unicodedata.combining(caracter)
    )

    return re.sub(
        r"[^A-Z0-9]+",
        "_",
        texto,
    ).strip("_")


def norm_codigo(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip().upper()

    if re.fullmatch(r"-?\d+\.0", texto):
        texto = texto[:-2]

    return re.sub(
        r"\s+",
        "",
        texto,
    )


def sha256(ruta):
    h = hashlib.sha256()

    with ruta.open("rb") as archivo:
        for bloque in iter(
            lambda: archivo.read(1024 * 1024),
            b"",
        ):
            h.update(bloque)

    return h.hexdigest()


def buscar_columna(df, candidatas, obligatoria=False):
    mapa = {
        norm_col(columna): columna
        for columna in df.columns
    }

    for candidata in candidatas:
        clave = norm_col(candidata)

        if clave in mapa:
            return mapa[clave]

    if obligatoria:
        raise RuntimeError(
            "No se encontró ninguna columna entre: "
            + " | ".join(candidatas)
        )

    return None


def unir_unicos(serie):
    valores = sorted({
        norm_codigo(valor)
        for valor in serie
        if norm_codigo(valor)
    })

    return " | ".join(valores)


def lista_numerica(serie):
    valores = pd.to_numeric(
        serie,
        errors="coerce",
    ).dropna()

    valores = sorted({
        float(valor)
        for valor in valores
    })

    salida = []

    for valor in valores:
        if float(valor).is_integer():
            salida.append(str(int(valor)))
        else:
            salida.append(str(valor))

    return " | ".join(salida)


if not FUENTE_EXCEL.exists():
    raise RuntimeError(
        f"No existe la fuente: {FUENTE_EXCEL}"
    )


libro = pd.ExcelFile(
    FUENTE_EXCEL,
    engine="openpyxl",
)

hojas_requeridas = {
    "DIAGNOSTICO_42",
    "DETALLE_MALLAS",
    "UNIVERSO_PLANES",
    "BLOQUEO_ORIGEN",
}

faltantes = hojas_requeridas - set(
    libro.sheet_names
)

if faltantes:
    raise RuntimeError(
        "Faltan hojas requeridas: "
        + " | ".join(sorted(faltantes))
    )


diagnostico = pd.read_excel(
    FUENTE_EXCEL,
    sheet_name="DIAGNOSTICO_42",
    dtype=str,
    keep_default_na=False,
    engine="openpyxl",
)

detalle = pd.read_excel(
    FUENTE_EXCEL,
    sheet_name="DETALLE_MALLAS",
    dtype=object,
    engine="openpyxl",
)

universo = pd.read_excel(
    FUENTE_EXCEL,
    sheet_name="UNIVERSO_PLANES",
    dtype=str,
    keep_default_na=False,
    engine="openpyxl",
)

bloqueo = pd.read_excel(
    FUENTE_EXCEL,
    sheet_name="BLOQUEO_ORIGEN",
    dtype=str,
    keep_default_na=False,
    engine="openpyxl",
)


if len(diagnostico) != 15:
    raise RuntimeError(
        f"Se esperaban 15 planes y se encontraron "
        f"{len(diagnostico)}."
    )

if diagnostico["PLAN_NORM"].duplicated().any():
    duplicados = diagnostico.loc[
        diagnostico["PLAN_NORM"].duplicated(
            keep=False
        ),
        "PLAN_NORM",
    ].tolist()

    raise RuntimeError(
        "Existen planes duplicados en el diagnóstico: "
        + " | ".join(sorted(set(duplicados)))
    )


col_plan_detalle = buscar_columna(
    detalle,
    ["PLAN_MALLA"],
    obligatoria=True,
)

col_nivel = buscar_columna(
    detalle,
    [
        "NIVEL_NUM",
        "NIVEL_ORIGINAL",
    ],
    obligatoria=True,
)

col_carrera = buscar_columna(
    detalle,
    ["CARRERA_MALLA"],
    obligatoria=False,
)

col_asignatura = buscar_columna(
    detalle,
    ["ASIGNATURA"],
    obligatoria=False,
)


detalle["PLAN_NORM"] = (
    detalle[col_plan_detalle]
    .map(norm_codigo)
)

detalle["NIVEL_ANALISIS"] = pd.to_numeric(
    detalle[col_nivel],
    errors="coerce",
)


conteo_nivel = (
    detalle.dropna(
        subset=["NIVEL_ANALISIS"]
    )
    .groupby(
        [
            "PLAN_NORM",
            "NIVEL_ANALISIS",
        ],
        dropna=False,
    )
    .agg(
        REGISTROS_NIVEL=(
            "PLAN_NORM",
            "size",
        ),
        CARRERAS_OBSERVADAS=(
            col_carrera,
            unir_unicos,
        )
        if col_carrera
        else (
            "PLAN_NORM",
            lambda serie: "",
        ),
        ASIGNATURAS_DISTINTAS=(
            col_asignatura,
            lambda serie: pd.Series(
                serie
            ).dropna().astype(str).str.strip().replace(
                "",
                pd.NA,
            ).dropna().nunique(),
        )
        if col_asignatura
        else (
            "PLAN_NORM",
            "size",
        ),
    )
    .reset_index()
)

conteo_nivel["NIVEL_ANALISIS"] = (
    conteo_nivel["NIVEL_ANALISIS"]
    .map(
        lambda valor: (
            int(valor)
            if float(valor).is_integer()
            else valor
        )
    )
)


resumen_detalle = (
    detalle.groupby(
        "PLAN_NORM",
        dropna=False,
    )
    .agg(
        NIVELES_DETECTADOS=(
            "NIVEL_ANALISIS",
            lista_numerica,
        ),
        NIVEL_MIN_DETECTADO=(
            "NIVEL_ANALISIS",
            "min",
        ),
        NIVEL_MAX_DETECTADO=(
            "NIVEL_ANALISIS",
            "max",
        ),
        CANTIDAD_NIVELES_DETECTADOS=(
            "NIVEL_ANALISIS",
            lambda serie: pd.to_numeric(
                serie,
                errors="coerce",
            ).dropna().nunique(),
        ),
        TOTAL_REGISTROS_MALLA=(
            "PLAN_NORM",
            "size",
        ),
    )
    .reset_index()
)


matriz = diagnostico.merge(
    resumen_detalle,
    on="PLAN_NORM",
    how="left",
    validate="one_to_one",
)


conteo_filas_bloqueo = (
    bloqueo.groupby(
        "PLAN_NORM",
        dropna=False,
    )
    .size()
    .reset_index(
        name="FILAS_PRECARGA_AFECTADAS"
    )
)

matriz = matriz.merge(
    conteo_filas_bloqueo,
    on="PLAN_NORM",
    how="left",
    validate="one_to_one",
)


columnas_contexto = [
    columna
    for columna in matriz.columns
    if columna.endswith("_VALORES")
]

columnas_base = [
    columna
    for columna in [
        "PLAN_NORM",
        "FILAS_PRECARGA_AFECTADAS",
        "CARRERAS_MALLA",
        "NIVELES_DETECTADOS",
        "NIVEL_MIN_DETECTADO",
        "NIVEL_MAX_DETECTADO",
        "CANTIDAD_NIVELES_DETECTADOS",
        "TOTAL_REGISTROS_MALLA",
        "ESTADO_EVIDENCIA_PERIODIZACION",
        "FUENTE_NECESARIA",
        "ACCION_REQUERIDA",
    ]
    if columna in matriz.columns
]

solicitud = matriz[
    columnas_base + columnas_contexto
].copy()


solicitud[
    "PREGUNTA_1_SEMANTICA_NIVEL"
] = (
    "¿Qué representa el campo NIVEL en este plan: "
    "semestre, trimestre, período académico, módulo, "
    "año curricular u otra unidad?"
)

solicitud[
    "RESPUESTA_1_SEMANTICA_NIVEL"
] = ""

solicitud[
    "PREGUNTA_2_CANTIDAD_ANIOS_PLAN"
] = (
    "¿Cuántos años curriculares se deben informar "
    "para este plan en Carreras Avance Curricular?"
)

solicitud[
    "RESPUESTA_2_CANTIDAD_ANIOS_PLAN"
] = ""

solicitud[
    "PREGUNTA_3_MAPA_NIVEL_ANIO"
] = (
    "Indique explícitamente qué niveles corresponden "
    "a cada año curricular."
)

solicitud[
    "ANIO_1_NIVELES"
] = ""

solicitud[
    "ANIO_2_NIVELES"
] = ""

solicitud[
    "ANIO_3_NIVELES"
] = ""

solicitud[
    "ANIO_4_NIVELES"
] = ""

solicitud[
    "ANIO_5_NIVELES"
] = ""

solicitud[
    "ANIO_6_NIVELES"
] = ""

solicitud[
    "ANIO_7_NIVELES"
] = ""

solicitud[
    "PREGUNTA_4_UNIDAD_MEDIDA"
] = (
    "Confirme la unidad de medida aplicable al plan "
    "y la distribución total de unidades por año."
)

solicitud[
    "TIPO_UNIDAD_MEDIDA_CONFIRMADA"
] = ""

solicitud[
    "TOTAL_UNIDADES_PLAN_CONFIRMADO"
] = ""

for anio in range(1, 8):
    solicitud[
        f"UNIDADES_ANIO_{anio}_CONFIRMADAS"
    ] = ""

solicitud[
    "FUENTE_DOCUMENTAL_RESPALDO"
] = ""

solicitud[
    "RUTA_O_NOMBRE_DOCUMENTO"
] = ""

solicitud[
    "RESPONSABLE_CONFIRMACION"
] = ""

solicitud[
    "UNIDAD_RESPONSABLE"
] = ""

solicitud[
    "FECHA_CONFIRMACION"
] = ""

solicitud[
    "OBSERVACIONES_INSTITUCIONALES"
] = ""

solicitud[
    "ESTADO_RESPUESTA"
] = "PENDIENTE"

solicitud[
    "USO_AUTORIZADO_EN_CARGA"
] = "NO"


resumen = pd.DataFrame([
    {
        "INDICADOR":
        "Filas físicas bloqueadas",
        "VALOR": len(bloqueo),
    },
    {
        "INDICADOR":
        "Planes únicos pendientes",
        "VALOR": len(solicitud),
    },
    {
        "INDICADOR":
        "Planes con solo NIVEL sin semántica",
        "VALOR": int(
            solicitud[
                "ESTADO_EVIDENCIA_PERIODIZACION"
            ].eq(
                "SOLO_NIVEL_SIN_SEMANTICA_CONFIRMADA"
            ).sum()
        ),
    },
    {
        "INDICADOR":
        "Conversiones aplicadas",
        "VALOR": 0,
    },
    {
        "INDICADOR":
        "Planes autorizados para carga",
        "VALOR": 0,
    },
])


instrucciones = pd.DataFrame([
    {
        "PASO": 1,
        "INSTRUCCION":
        "Completar una fila por cada PLAN_NORM.",
    },
    {
        "PASO": 2,
        "INSTRUCCION":
        "Definir qué representa NIVEL en ese plan.",
    },
    {
        "PASO": 3,
        "INSTRUCCION":
        "Indicar los niveles que pertenecen a cada año curricular.",
    },
    {
        "PASO": 4,
        "INSTRUCCION":
        "Confirmar la unidad de medida y las unidades por año.",
    },
    {
        "PASO": 5,
        "INSTRUCCION":
        "Adjuntar o identificar la fuente documental utilizada.",
    },
    {
        "PASO": 6,
        "INSTRUCCION":
        "Registrar responsable, unidad y fecha de confirmación.",
    },
    {
        "PASO": 7,
        "INSTRUCCION":
        "No modificar PLAN_NORM, niveles detectados ni conteos.",
    },
    {
        "PASO": 8,
        "INSTRUCCION":
        "Marcar ESTADO_RESPUESTA como CONFIRMADO solo "
        "cuando todos los campos estén respaldados.",
    },
])


solicitud.to_csv(
    RESULTADOS
    / "01_SOLICITUD_PERIODIZACION_15_PLANES.tsv",
    sep="\t",
    index=False,
)

conteo_nivel.to_csv(
    RESULTADOS
    / "02_DISTRIBUCION_OBSERVADA_POR_NIVEL.tsv",
    sep="\t",
    index=False,
)

resumen.to_csv(
    AUDITORIA
    / "01_RESUMEN.tsv",
    sep="\t",
    index=False,
)


excel = (
    ENTREGA
    / "SOLICITUD_CONFIRMACION_PERIODIZACION_15_PLANES.xlsx"
)

with pd.ExcelWriter(
    excel,
    engine="openpyxl",
) as writer:
    resumen.to_excel(
        writer,
        sheet_name="RESUMEN",
        index=False,
    )

    instrucciones.to_excel(
        writer,
        sheet_name="INSTRUCCIONES",
        index=False,
    )

    solicitud.to_excel(
        writer,
        sheet_name="CONFIRMAR_15_PLANES",
        index=False,
    )

    conteo_nivel.to_excel(
        writer,
        sheet_name="DETALLE_POR_NIVEL",
        index=False,
    )

    bloqueo.to_excel(
        writer,
        sheet_name="42_FILAS_AFECTADAS",
        index=False,
    )

    detalle.to_excel(
        writer,
        sheet_name="DETALLE_MALLAS",
        index=False,
    )

    universo.to_excel(
        writer,
        sheet_name="UNIVERSO_PLANES_ORIGEN",
        index=False,
    )

    for hoja in writer.book.worksheets:
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        hoja.sheet_view.showGridLines = False

        for celda in hoja[1]:
            celda.font = Font(
                bold=True,
                color="FFFFFF",
            )

            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for columna in hoja.columns:
            letra = get_column_letter(
                columna[0].column
            )

            ancho = max(
                len(str(celda.value or ""))
                for celda in columna[:2000]
            )

            hoja.column_dimensions[letra].width = min(
                max(ancho + 2, 12),
                55,
            )

        for fila in hoja.iter_rows(
            min_row=2
        ):
            for celda in fila:
                celda.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )


hash_fuente = sha256(FUENTE_EXCEL)
hash_salida = sha256(excel)

hashes = pd.DataFrame([
    {
        "TIPO": "FUENTE",
        "ARCHIVO": str(FUENTE_EXCEL),
        "SHA256": hash_fuente,
    },
    {
        "TIPO": "SALIDA",
        "ARCHIVO": str(excel),
        "SHA256": hash_salida,
    },
])

hashes.to_csv(
    AUDITORIA
    / "02_HASHES.tsv",
    sep="\t",
    index=False,
)


validaciones = pd.DataFrame([
    {
        "VALIDACION": "PLANES_UNICOS",
        "RESULTADO":
        "OK" if len(solicitud) == 15 else "ERROR",
        "DETALLE": len(solicitud),
    },
    {
        "VALIDACION": "PLANES_DUPLICADOS",
        "RESULTADO":
        "OK"
        if not solicitud["PLAN_NORM"].duplicated().any()
        else "ERROR",
        "DETALLE": int(
            solicitud["PLAN_NORM"]
            .duplicated()
            .sum()
        ),
    },
    {
        "VALIDACION": "FILAS_BLOQUEO_CONSERVADAS",
        "RESULTADO":
        "OK" if len(bloqueo) == 42 else "ERROR",
        "DETALLE": len(bloqueo),
    },
    {
        "VALIDACION": "CONVERSIONES_APLICADAS",
        "RESULTADO": "OK",
        "DETALLE": 0,
    },
    {
        "VALIDACION": "FASE_4_DESBLOQUEADA",
        "RESULTADO": "NO",
        "DETALLE":
        "Pendiente confirmación institucional de 15 planes.",
    },
])

validaciones.to_csv(
    AUDITORIA
    / "03_VALIDACIONES.tsv",
    sep="\t",
    index=False,
)


manifiesto = {
    "fecha_ejecucion": datetime.now().isoformat(),
    "proceso": "Avance Curricular SIES 2026",
    "subproceso": "Carreras Avance Curricular",
    "fase": "4B",
    "filas_fisicas_bloqueadas": len(bloqueo),
    "planes_unicos_pendientes": len(solicitud),
    "estado_evidencia": (
        "SOLO_NIVEL_SIN_SEMANTICA_CONFIRMADA"
    ),
    "conversiones_aplicadas": 0,
    "fuentes_modificadas": False,
    "fase_4_desbloqueada": False,
    "archivo_solicitud": str(excel),
    "sha256_fuente": hash_fuente,
    "sha256_salida": hash_salida,
}

(SALIDA / "manifest_ejecucion.json").write_text(
    json.dumps(
        manifiesto,
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)


print()
print("=" * 108)
print("SOLICITUD INSTITUCIONAL — PERIODIZACIÓN DE 15 PLANES")
print("=" * 108)
print(f"Filas físicas afectadas: {len(bloqueo)}")
print(f"Planes únicos por confirmar: {len(solicitud)}")
print(
    "Planes con solo NIVEL sin semántica confirmada: "
    f"{solicitud['ESTADO_EVIDENCIA_PERIODIZACION'].eq(
        'SOLO_NIVEL_SIN_SEMANTICA_CONFIRMADA'
    ).sum()}"
)
print()
print(f"Excel institucional: {excel}")
print(f"TSV de respaldo: {RESULTADOS / '01_SOLICITUD_PERIODIZACION_15_PLANES.tsv'}")
print(f"Detalle por nivel: {RESULTADOS / '02_DISTRIBUCION_OBSERVADA_POR_NIVEL.tsv'}")
print(f"Carpeta: {SALIDA}")
print("Conversiones aplicadas: 0")
print("Fuentes modificadas: NO")
print("Fase 4 desbloqueada: NO")
print("=" * 108)
