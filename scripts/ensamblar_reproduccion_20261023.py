"""Inyecta hojas de auditoría creadas con Artifact Tool; conserva todas las hojas originales byte a byte."""
import zipfile,re,json,hashlib,sys,shutil,xml.etree.ElementTree as E
from pathlib import Path
from xml.sax.saxutils import escape
import openpyxl
O=Path(sys.argv[1]);R=Path(__file__).resolve().parents[1]
src=R/'outputs/retencion_puente_provisional_20261023_v10/RETENCION_IPSS_PROVISIONAL_PUENTE_20261023_V10.xlsx'
dst=O/'RETENCION_IPSS_PUENTE_REPRODUCIDO_20261023_V11.xlsx'
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
assert sha(src)=='d148c60042df2e69ba80d4420074518cb91045686b1c281f977093f035a2bd32'
z=zipfile.ZipFile(src);new=zipfile.ZipFile(O/'suplemento.xlsx');payload=json.load(open(O/'payload_agregado.json'))
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda x:'{'+NS+'}'+x;E.register_namespace('',NS)
sn=[] if 'xl/sharedStrings.xml' not in new.namelist() else [''.join(n.itertext()) for n in E.fromstring(new.read('xl/sharedStrings.xml'))]
base=z.read('xl/styles.xml').decode();b=E.fromstring(base);a=E.fromstring(new.read('xl/styles.xml'));offset={};num_map={}
for tag in ['numFmts','fonts','fills','borders','cellXfs']:
 old=b.find(q(tag));extra=a.find(q(tag));offset[tag]=len(old) if old is not None else 0
 if extra is None:continue
 if tag=='numFmts':
  nextid=max([163]+[int(n.get('numFmtId')) for n in old] if old is not None else [163])+1
  for n in extra:num_map[int(n.get('numFmtId'))]=nextid;n.set('numFmtId',str(nextid));nextid+=1
 if tag=='cellXfs':
  for n in extra:
   for attr,t in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:n.set(attr,str(int(n.get(attr,'0'))+offset[t]))
   n.set('numFmtId',str(num_map.get(int(n.get('numFmtId','0')),int(n.get('numFmtId','0')))));n.set('xfId','0')
 children=''.join(E.tostring(n,encoding='unicode').replace(' xmlns="'+NS+'"','') for n in extra)
 found=re.search(rf'<{tag}\b[^>]*>.*?</{tag}>',base,re.S)
 if found:
  chunk=found.group();chunk=re.sub(r'count="\d+"',f'count="{offset[tag]+len(extra)}"',chunk,count=1);chunk=chunk.replace(f'</{tag}>',children+f'</{tag}>');base=base[:found.start()]+chunk+base[found.end():]
 else:base=base.replace('<fonts',f'<{tag} count="{len(extra)}">{children}</{tag}><fonts',1)
def clean(txt):
 txt=txt.replace('<x:','<').replace('</x:','</').replace('xmlns:x=','xmlns=')
 def cell(m):
  s=m.group()
  if re.search(r'\bt="s"',s):
   idx=int(re.search(r'<v>(\d+)</v>',s).group(1));s=re.sub(r'\bt="s"','t="inlineStr"',s);s=re.sub(r'<v>.*?</v>','<is><t xml:space="preserve">'+escape(sn[idx])+'</t></is>',s,flags=re.S)
  if re.search(r'\bs="\d+"',s):s=re.sub(r'\bs="(\d+)"',lambda m:f's="{int(m.group(1))+offset["cellXfs"]}"',s,count=1)
  else:s=s.replace('<c ','<c s="'+str(offset['cellXfs'])+'" ',1)
  return s
 return re.sub(r'<c\b[^>]*(?<!/)>.*?</c>',cell,txt,flags=re.S)
work=z.read('xl/workbook.xml').decode();rels=z.read('xl/_rels/workbook.xml.rels').decode();ct=z.read('[Content_Types].xml').decode()
oldnames=[s.get('name') for s in E.fromstring(work).find(q('sheets'))]
assert len(oldnames)==48
repl={}
for i,p in enumerate(payload,1):
 assert p['name'] not in oldnames
 j=len(oldnames)+i;part=f'xl/worksheets/sheet{j}.xml';assert part not in z.namelist()
 repl[part]=clean(new.read(f'xl/worksheets/sheet{i}.xml').decode())
 work=work.replace('</sheets>',f'<sheet name="{escape(p["name"])}" sheetId="{j}" r:id="rIdAudit{j}"/></sheets>')
 rels=rels.replace('</Relationships>',f'<Relationship Id="rIdAudit{j}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{j}.xml"/></Relationships>')
 ct=ct.replace('</Types>',f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
repl.update({'xl/styles.xml':base,'xl/workbook.xml':work,'xl/_rels/workbook.xml.rels':rels,'[Content_Types].xml':ct})
with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
 for name in z.namelist():out.writestr(name,repl.pop(name).encode() if name in repl else z.read(name))
 for name,data in repl.items():out.writestr(name,data.encode())
zz=zipfile.ZipFile(dst);assert zz.testzip() is None
changed=[n for n in z.namelist() if z.read(n)!=zz.read(n)]
assert set(changed)=={'xl/styles.xml','xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml'}
charts=[n for n in z.namelist() if re.fullmatch('xl/charts/chart[0-9]+.xml',n)]
assert len(charts)==13 and all(z.read(n)==zz.read(n) for n in charts)
for part in zz.namelist():
 if part.endswith('.xml') or part.endswith('.rels'):E.fromstring(zz.read(part))
w=openpyxl.load_workbook(dst,read_only=True,data_only=True);wf=openpyxl.load_workbook(dst,read_only=True,data_only=False)
s=w['49_Reproduccion_V10'];assert s['B8'].value==5901==s['C8'].value and s['B9'].value==1838 and s['C9'].value==1751
assert abs(s['C10'].value-1751/5901)<1e-12
assert s['B16'].value==s['C16'].value==0
s=w['50_Prueba_Personas'];assert s['B11'].value==s['C11'].value==17826 and s['C6'].value==1387
s=w['51_Evidencias_Releidas'];assert s['B38'].value==s['C38'].value==49218
errors=[];nf=0;cachemissing=[]
for p in payload:
 s=w[p['name']]
 for c,formula in p['formulas'].items():
  assert wf[p['name']][c].value==formula;nf+=1
  if s[c].value is None:cachemissing.append([p['name'],c])
 for row in s:
  for c in row:
   if c.data_type=='e':errors.append([s.title,c.coordinate,c.value])
assert not errors and not cachemissing
qa={'estado':'OK_CON_ADVERTENCIAS','hojas':len(w.sheetnames),'hojas_originales_bytes_identicos':48,'graficos_bytes_identicos':13,
 'partes_originales_modificadas':changed,'formulas_nuevas_cache_verificado':nf,'errores_hojas_nuevas':errors,'fuente_sha256':sha(src),
 'libro_sha256':sha(dst),'limites':['Sin validación de apertura en Excel nativo; recálculo Artifact Tool y lectura independiente openpyxl',
 'Hojas históricas preservadas: no se certifican globalmente sus tasas; desfase conocido hoja 31 se conserva y advierte']}
w.close();wf.close()
desktop=Path.home()/'Desktop'/dst.name
if desktop.exists():assert (O/'QA_EXCEL.json').exists() and sha(desktop)==json.load(open(O/'QA_EXCEL.json'))['copia_escritorio_sha256']
shutil.copy2(dst,desktop);assert sha(dst)==sha(desktop)
qa['copia_escritorio']=str(desktop);qa['copia_escritorio_sha256']=sha(desktop)
(O/'QA_EXCEL.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2));print(json.dumps(qa,ensure_ascii=False))
