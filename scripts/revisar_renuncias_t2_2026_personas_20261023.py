"""Contraste agregado de renuncias T2; nunca exporta identificadores personales.

Lectura exclusiva de fuentes. No adjudica renuncias mediante estados posteriores.
Salida en carpeta nueva; no modifica libros, denominadores ni motores.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import json
import subprocess
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CAR = ROOT.parent / 'CARACTERIZACIÓN'
FROZEN = ROOT / '.local_restricted/retencion_intermedia_t2_t3/20261022_fuentes'
OUT = ROOT / 'outputs/renuncias_t2_2026_personas_20261023'


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def norm(value):
    return '' if value is None else str(value).strip().upper()


def rows(path):
    book = load_workbook(path, read_only=True, data_only=True)
    it = book['Hoja1'].iter_rows(values_only=True)
    header = next(it)
    result = [dict(zip(header, row)) for row in it if any(x is not None for x in row)]
    book.close()
    return result


def git(*args):
    return subprocess.check_output(['git', '-C', str(ROOT), *args], text=True).strip()


def main():
    if OUT.exists():
        raise SystemExit('La salida ya existe; conservarla y usar una nueva versión.')
    source = next(FROZEN.glob('*T_2_13_07_2026.xlsx'))
    t3 = next(FROZEN.glob('*T3_02_10_2026.xlsx'))
    sources = {
        'T2': source,
        'T3': t3,
        'DA10246': CAR / '.local/fuentes_congeladas/2026/P2/2026_P2_20260721_161905/DATOSDEALUMNOS_10246.xlsx',
        'DA11998': CAR / '.local/analisis_anual/2026/20261016_002933_892820/entrega/03_RESPALDOS/P3/01_ENTRADAS/DATOSDEALUMNOS_11998.xlsx',
        'DA12159': Path('/Users/alexi/Desktop/Caracterizacion_2026-10-22_2027_P1_084218/01_ENTRADAS/DATOSDEALUMNOS_12159.xlsx'),
        'MANIFEST_ORIGEN': ROOT / 'procesos/retencion_intermedia_t2_t3/MANIFEST_RETENCION_FOTO_INTERTRIMESTRAL_20261022.json',
        'MANIFEST_10246': CAR / 'docs/gobernanza/FUENTES_CONGELADAS_2026_P2.json',
        'MANIFEST_12159': Path('/Users/alexi/Desktop/Caracterizacion_2026-10-22_2027_P1_084218/04_AUDITORIA/manifest_ejecucion.json'),
        'V16': ROOT / 'outputs/aclaracion_tasas_2026_20261023_v16/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V16.xlsx',
        'V16_ESCRITORIO': Path('/Users/alexi/Desktop/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V16.xlsx'),
    }
    before = {k: digest(v) for k, v in sources.items()}
    head, branch = git('rev-parse', 'HEAD'), git('branch', '--show-current')
    population = rows(source)
    assert len(population) == 528
    assert len({norm(r['CODCLI']) for r in population}) == 528
    person_counts = Counter(norm(r['RUT']) for r in population)
    assert '' not in person_counts and len(person_counts) == 527
    assert Counter(person_counts.values()) == Counter({1: 526, 2: 1})
    assert all(norm(r['ANOMATRICULA']) == '2026' and norm(r['PERIODOMATRICULA']) == '2'
               and norm(r['ANOINGRESO']) == '2026' and norm(r['PERIODOINGRESO']) == '2' for r in population)

    book = load_workbook(source, read_only=True, data_only=True)
    original = [{'carrera': r[0], 'matriculas': r[1], 'renuncias_celda': r[2],
                 'renuncias_calculo': r[2] or 0, 'fila': i}
                for i, r in enumerate(book['Hoja2'].iter_rows(min_row=32, max_row=50, values_only=True), 32)]
    sheets = [{'nombre': s.title, 'filas': s.max_row, 'columnas': s.max_column,
               'visibilidad': s.sheet_state} for s in book]
    book.close()
    assert sum(r['matriculas'] for r in original) == 528
    assert sum(r['renuncias_calculo'] for r in original) == 66
    book = load_workbook(source, read_only=True, data_only=False)
    assert all(book['Hoja2'][f'C{i}'].data_type != 'f' for i in range(32, 51))
    book.close()
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(source) as z:
        cache = ET.fromstring(z.read('xl/pivotCache/pivotCacheDefinition1.xml'))
        cache_info = {'registros': int(cache.attrib['recordCount']),
                      'origen': cache.find('s:cacheSource/s:worksheetSource', ns).attrib,
                      'campos': [e.attrib['name'] for e in cache.find('s:cacheFields', ns)]}
    assert cache_info['registros'] == 528

    comparisons = {}
    for key in ['DA10246', 'DA11998', 'DA12159']:
        data = rows(sources[key])
        index, by_rut = defaultdict(list), defaultdict(list)
        for row in data:
            index[norm(row['CODCLI'])].append(row)
            by_rut[norm(row['RUT'])].append(row)
        states, situations, careers, cardinality = Counter(), Counter(), defaultdict(Counter), Counter()
        duplicate_states, unmatched_research = [], []
        rut_equal = 0
        alive_persons = set()
        for row in population:
            match = index[norm(row['CODCLI'])]
            cardinality[str(len(match))] += 1
            if len(match) == 1:
                target = match[0]
                equal = norm(row['RUT']) == norm(target['RUT'])
                assert equal, 'Identidad discordante: no continuar unión.'
                rut_equal += int(equal)
                state, situation = norm(target['ESTADOACADEMICO']), norm(target['SITUACION'])
                situations[state + ' | ' + situation] += 1
                if state == 'VIGENTE':
                    alive_persons.add(norm(row['RUT']))
                if person_counts[norm(row['RUT'])] == 2:
                    duplicate_states.append({'carrera': row['CODCARPR'], 'estado': state, 'situacion': situation})
            elif len(match) == 0:
                state = 'NO_LOCALIZADO_POR_CODCLI'
                other = by_rut[norm(row['RUT'])]
                unmatched_research.append({'carrera_origen': row['CODCARPR'], 'filas_mismo_rut': len(other),
                    'filas_misma_carrera': sum(norm(x['CODCARPR']) == norm(row['CODCARPR']) for x in other),
                    'estados_otro_codcli': dict(Counter(norm(x['ESTADOACADEMICO']) for x in other))})
            else:
                raise AssertionError('CODCLI no único en respaldo: no usar unión muchos-a-muchos.')
            states[state] += 1
            careers[row['NOMBRE_L']][state] += 1
        per_career = []
        for item in original:
            count = careers[item['carrera']]
            assert sum(count.values()) == item['matriculas']
            # Solo un control aritmético: estas categorías NO son una definición de renuncia.
            complement = sum(v for s, v in count.items() if s != 'VIGENTE')
            per_career.append({**item, 'estados': dict(count),
                'complemento_vigente_incluye_no_localizados': complement,
                'diferencia_complemento_vs_renuncias': complement - item['renuncias_calculo']})
        comparisons[key] = {'clasificacion': 'B/C', 'hoja': 'Hoja1',
            'columnas_cruce': ['CODCLI', 'RUT'], 'columnas_estado': ['ESTADOACADEMICO', 'SITUACION'],
            'filas_respaldo': len(data), 'cardinalidad_codcli': dict(cardinality),
            'coincidencias_rut': rut_equal, 'estados': dict(states), 'situaciones': dict(situations),
            'personas_con_matricula_vigente_en_respaldo': len(alive_persons),
            'duplicado_sin_identificadores': duplicate_states, 'busqueda_rut_sin_codcli': unmatched_research,
            'por_carrera': per_career,
            'carreras_discordantes_complemento': sum(x['diferencia_complemento_vs_renuncias'] != 0 for x in per_career),
            'limite': 'Estado del respaldo; no acredita pertenencia a las66 renuncias del cierreT2.'}
        assert sum(states.values()) == 528

    scenarios = [
        {'condicion': 'Ambas matrículas de la única persona duplicada permanecen después de restar las66.',
         'personas_restantes': 461, 'personas_base_escenario': 527, 'tasa': 461 / 527},
        {'condicion': 'Una o ambas matrículas de esa persona forman parte de las66.',
         'personas_restantes': 462, 'personas_base_escenario': 527, 'tasa': 462 / 527},
    ]
    result = {'proceso': 'Retención intertrimestral interna IPSS162', 'periodo': '2026/T2',
        'fecha_ejecucion': datetime.now().astimezone().isoformat(),
        'clasificacion_evidencia': 'B observado; C cálculo; E atribución pendiente',
        'actual_sin_cambios': {'matriculas': 528, 'personas': 527, 'renuncias_agregadas': 66,
                             'matriculas_restantes': 462, 'tasa_matriculas': 462 / 528},
        'estructura_t2': sheets, 'cache_dinamica_t2': cache_info,
        'detalle_renuncias': {'estado': 'NO_LOCALIZADO_EN_COBERTURA_REVISADA', 'clasificacion': 'E',
          'referencia': 'T2 Hoja2 C32:C50: constantes agregadas, sin lista individual acreditada.'},
        'contrastes': comparisons, 'escenarios_condicionales_c_e': scenarios,
        'nota_escenario': '527 es unidad personas del mismo universo528; no reemplaza ningún denominador aprobado.',
        'supuestos_escenario_e': [
            'Las66 renuncias representan66 matrículas distintas del universo528, sin repeticiones ni otros períodos.',
            'Una persona permanece si conserva al menos una matrícula de ese universo.',
            'Restar528−66 reproduce la regla del resumen; no acredita vigencia individual al cierre.'
        ],
        'fuentes': {k: {'ruta': str(v), 'sha256': before[k]} for k, v in sources.items()},
    }
    after = {k: digest(v) for k, v in sources.items()}
    assert before == after
    assert head == git('rev-parse', 'HEAD') and branch == git('branch', '--show-current')
    assert before['V16'] == before['V16_ESCRITORIO']
    qa = {'integridad_fuentes': 'OK', 'identidad_codcli_rut_sin_discordancias': 'OK',
        'control_528_matriculas_527_personas_66_renuncias': 'OK',
        'sin_uniones_muchos_a_muchos': 'OK', 'conciliacion_estados_528_por_respaldo': 'OK',
        'V16_y_escritorio_sin_cambios': 'OK', 'identificadores_exportados': False,
        'tasa_por_personas_del_cierre_T2': 'PENDIENTE: falta atribuir la nómina66',
        'rama_antes_despues': branch, 'head_antes_despues': head, 'commit_push': False}
    OUT.mkdir(parents=True)
    for name, value in [('EVIDENCIA_AGREGADA.json', result), ('QA.json', qa)]:
        (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'salida': str(OUT), 'control_actual': result['actual_sin_cambios'],
        'estados_por_respaldo': {k: v['estados'] for k, v in comparisons.items()},
        'carreras_discordantes_12159': comparisons['DA12159']['carreras_discordantes_complemento'],
        'escenarios_condicionales': scenarios, 'qa': qa}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
