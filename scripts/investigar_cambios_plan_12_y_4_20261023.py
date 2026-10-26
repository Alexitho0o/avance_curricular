"""Relectura dirigida de fuentes: distingue código anual, etiqueta de plan y fecha efectiva.
Solo genera derivados; ningún candidato se adopta ni se modifica una fuente.
"""
import json,csv,re,hashlib,datetime,sys,subprocess
from pathlib import Path
from collections import defaultdict,Counter
import openpyxl
R=Path(__file__).resolve().parents[1];cfg=json.load(open(sys.argv[1]));O=Path(cfg['output']);P=Path(cfg['privado'])
assert subprocess.run(['git','-C',str(P),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
load=lambda p:json.load(open(p));sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str))
def val(x):return '' if x is None else str(int(x)) if isinstance(x,float) and x.is_integer() else str(x).strip()
def norm(x,dv=None):
 s=val(x).upper().replace('.','').replace(' ','')
 if dv is not None:s+='-'+val(dv).upper()
 if not re.fullmatch(r'\d+-[0-9K]',s):return None
 a,b=s.split('-');return str(int(a))+'-'+b
def date(x):
 if isinstance(x,datetime.datetime):return x.date().isoformat()
 if isinstance(x,datetime.date):return x.isoformat()
 for fmt in ['%Y-%m-%d %H:%M:%S','%Y-%m-%d','%d-%m-%Y %H:%M','%d-%m-%Y %H:%M:%S','%d/%m/%Y','%d/%m/%Y %H:%M:%S']:
  try:return datetime.datetime.strptime(val(x),fmt).date().isoformat()
  except ValueError:pass
 return None
F=load(P/'FOCO_PRIVADO.json');people={x['persona'] for x in F};assert len(F)==len(people)==16
B=R/'outputs/trayectorias_374_20261023_072401';BM=load(B/'MANIFEST.json');BP=Path(BM['privado']);hashes={};inventory=[];findings=[]
for n in ['CAMBIOS_VERSION_ANUAL.json','CUATRO_PISTAS_FECHA.json','OBSERVACIONES_REVISADAS.json','SEGMENTOS_EXPEDIENTE_PLAN.json']:
 assert sha(BP/n)==BM['salidas_privadas_hashes'][n];hashes[str(BP/n)]=sha(BP/n)
prior=[x for x in load(BP/'OBSERVACIONES_REVISADAS.json') if x['persona'] in people]
segments=[x for x in load(BP/'SEGMENTOS_EXPEDIENTE_PLAN.json') if x['persona'] in people]
# Sources are deliberately explicit, never "latest".
prof=load(P/'PERFIL_FUENTES_ADICIONALES.json');paths=P.joinpath('INVENTARIO_ONEDRIVE_RUTAS.txt').read_text().splitlines()
specs=[]
for x in prof:
 if Path(x['path']).name.startswith('PROMEDIOSDEALUMNOS_') and x['sheet']=='Hoja1':specs.append((x['path'],x['sheet'],1,'PROMEDIOS_ADICIONAL'))
 if 'CUMPLIMIENTO' in Path(x['path']).name.upper() or 'Cumplimiento Ingreso' in x['path']:
  if x['headers'] and any('CODIGO_PLAN_ESTUDIO' in h['columns'] for h in x['headers']):specs.append((x['path'],x['sheet'],x['headers'][0]['row'],'CALIFICACIONES_ASIGNATURA'))
 if Path(x['path']).name=='DATOSDEALUMNOS ORIGEN DE LOS DATOS2025.xlsx' and x['sheet']=='Hoja1':specs.append((x['path'],x['sheet'],1,'DATOSALUMNOS_HISTORICO'))
# Earlier PROMEDIOS, frozen 2026 P2; compare hash before accepting as independent snapshot.
p2=R.parent/'CARACTERIZACIÓN/.local/fuentes_congeladas/2026/P2/2026_P2_20260721_161905/PROMEDIOSDEALUMNOS_7804.xlsx'
if p2.is_file():specs.append((str(p2),'Hoja1',1,'PROMEDIOS_RESPALDO_P2'))
for s in paths:
 if Path(s).name=='DATOSDEALUMNOS_1758_CRUDO_10_04_2025.xlsx':specs.append((s,'Hoja1',1,'DATOSALUMNOS_HISTORICO'))
known={};rows=[];errors=[]
for path,sheet,hr,role in specs:
 try:
  h=sha(path);hashes[path]=h
  if (h,sheet) in known:inventory.append({'archivo':path,'sha256':h,'hoja':sheet,'estado':'DUPLICADO_BYTES','alias_de':known[h,sheet]});continue
  known[h,sheet]=path
  w=openpyxl.load_workbook(path,read_only=True,data_only=True)
  if sheet not in w.sheetnames:
   inventory.append({'archivo':path,'sha256':h,'estado':'HOJA_NO_LOCALIZADA','hojas':w.sheetnames});w.close();continue
  ws=w[sheet];it=ws.iter_rows(min_row=hr,values_only=True);heads=[val(x).upper() for x in next(it)];n=found=0
  # Repeated headers kept with explicit suffix; no accidental last-column substitution.
  counts=Counter();hh=[]
  for v in heads:counts[v]+=1;hh.append(v if counts[v]==1 else v+'__'+str(counts[v]))
  for rn,values in enumerate(it,hr+1):
   n+=1;d=dict(zip(hh,values));raw=d.get('RUT');k=norm(raw) or norm(raw,d.get('DIG'))
   if k not in people:continue
   keep=['CODCLI','CODCARPR','CODCARR','CODCARRERA_ALUMNO','JORNADA','SEDE','REGIMEN','PLAN_DE_ESTUDIO','CODIGO_PLAN_ESTUDIO','ANO','PERIODO','PERIODO__2','ANOMATRICULA','PERIODOMATRICULA','ANOINGRESO','PERIODOINGRESO','FECHAMATRICULA','FECHAFIRMA','CODRAMO','RAMOEQUIV','CODIGO_ASIGNATURA','SECCION','FECHAINICIO','FECHAFIN','FECHA REPORTE','FECHA MEDICION','ESTADO','ESTADO_ACADEMICO','ESTADOACADEMICO','SITUACION','DESCRIPCION_ESTADO','CONVALIDADO']
   z={t:val(d.get(t)) for t in keep if t in d};z.update(persona=k,RUT_original=val(raw),DV_original=val(d.get('DIG')),archivo=path,sha256=h,hoja=sheet,fila=rn,rol=role)
   z['fechas_interpretadas']={t:date(d.get(t)) for t in ['FECHAMATRICULA','FECHAFIRMA','FECHAINICIO','FECHAFIN','FECHA REPORTE','FECHA MEDICION'] if t in d}
   rows.append(z);found+=1
  inventory.append({'archivo':path,'sha256':h,'hoja':sheet,'fila_encabezado':hr,'encabezados':hh,'filas_leidas':n,'filas_foco':found,'rol':role,'corte_efectivo':'NO_INFERIDO_DEL_NOMBRE','estado':'LEIDO'});w.close()
  print(json.dumps({'fuente':Path(path).name,'hoja':sheet,'filas':n,'foco':found}),flush=True)
 except (OSError,ValueError,KeyError) as e:errors.append({'archivo':path,'hoja':sheet,'error':str(e)});print('Lectura pendiente: '+Path(path).name,flush=True)
dump(P/'FILAS_ADICIONALES_PRIVADAS.json',rows);dump(O/'INVENTARIO_ADICIONAL.json',inventory);dump(O/'ERRORES_LECTURA.json',errors)
# Re-read actual annual RAW sources used by the consolidation, including non-consolidated fields.
root=next(Path('/Users/alexi/Library/CloudStorage').glob('OneDrive-Instituto*'))
base=root/'Datos DAI (RAW) - 1.1 Datos crudos/1.2.1 Datos Institucionales/1.2.1.1 SIES/Enviados/Matrícula Unificada'
mu=[]
for year in [2022,2023,2024,2025]:
 if year in [2022,2023]:path=base/str(year)/(str(year)+'.csv');sheet='CSV';hr=1
 elif year==2024:path=next((base/'2024').glob('08-05-2024*Pregrado*csv'));sheet='CSV';hr=1
 else:path=base/'2025/2025.xlsx';sheet='08-05-2025_22-26-22_Matrícula_P';hr=1
 hashes[str(path)]=sha(path)
 if sheet=='CSV':
  raw=path.read_bytes()
  for enc in ['utf-8-sig','cp1252']:
   try:text=raw.decode(enc);break
   except UnicodeDecodeError:pass
  rr=csv.DictReader(text.splitlines(),delimiter=';');it=enumerate(rr,2);heads=rr.fieldnames
 else:
  w=openpyxl.load_workbook(path,read_only=True,data_only=True);ws=w[sheet];ii=iter(ws.values);heads=next(ii);it=((nr,dict(zip(heads,v))) for nr,v in enumerate(ii,2))
 n=found=0
 for rn,d in it:
  n+=1;k=norm(d.get('N_DOC'),d.get('DV'))
  if k not in people:continue
  code='I162S%sC%sJ%sV%s'%tuple(val(d.get(t)) for t in ['COD_SED','COD_CAR','JOR','VERSION'])
  z={'persona':k,'ano':year,'codigo':code,'familia':code.rsplit('V',1)[0],'archivo':str(path),'hoja':sheet,'fila':rn,'sha256':hashes[str(path)]}
  z.update({t:val(d.get(t)) for t in ['TIPO_DOC','N_DOC','DV','COD_SED','COD_CAR','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FECHA_MATRICULA','FECH_CAR','REINCORPORACION']});mu.append(z);found+=1
 inventory.append({'archivo':str(path),'sha256':hashes[str(path)],'hoja':sheet,'filas_leidas':n,'filas_foco':found,'encabezados':list(heads),'rol':'MU_PRIMARIA_SEGUN_CONSTRUCTOR','estado':'LEIDO'})
 if sheet!='CSV':w.close()
 print(json.dumps({'fuente':path.name,'filas':n,'foco':found}),flush=True)
dump(P/'MU_PRIMARIA_PRIVADA.json',mu)
# Main diagnosis: source lineage plus plan associations, without pretending subject dates are plan validity.
caseout=[]
for case in F:
 person=case['persona'];fam=case['familia'];r=[x for x in rows if x['persona']==person];ss=[x for x in segments if x['persona']==person and x['familia']==fam];mm=[x for x in mu if x['persona']==person and x['familia']==fam]
 pr=[x for x in prior if x['persona']==person and fam in x['familias']]
 planmatches=[];conflicts=[]
 for x in r:
  if x['rol']!='CALIFICACIONES_ASIGNATURA':continue
  contexts=[s for s in ss if s['CODCLI']==x.get('CODCLI') and s['carrera']==x.get('CODCARRERA_ALUMNO')]
  match=[s for s in contexts if s['plan']==x.get('CODIGO_PLAN_ESTUDIO')]
  z={'fila_asignatura':x,'planes_en_registro_previo':[s['plan'] for s in contexts],'coincide_plan':bool(match)}
  if match:planmatches.append(z)
  elif contexts:conflicts.append(z)
 primarychecks=[]
 for x in pr:
  if int(x['ano']) not in [2022,2023,2024,2025]:continue
  derived=set(x['codigos_anuales']);rawcodes={m['codigo'] for m in mm if m['ano']==int(x['ano'])}
  primarychecks.append({'observacion_id':x['observacion_id'],'ano':x['ano'],'codigos_derivado':sorted(derived),'codigos_raw':sorted(rawcodes),'coinciden':derived==rawcodes})
 # Notes may overlap years, but are not a dated assignment history.
 original=case['detalle_anterior'];start=original.get('ultimo_ano_codigo_anterior');end=original.get('primer_ano_codigo_siguiente')
 sameplans=[]
 if start and end:
  for s in ss:
   years={int(t[0]) for t in s['periodos_observados']}
   if start in years and end in years:sameplans.append({'CODCLI':s['CODCLI'],'carrera':s['carrera'],'plan':s['plan']})
 z={'caso':case['caso'],'tipo':case['tipo'],'persona':person,'familia':fam,'antecedente':original,'segmentos_previos':ss,'filas_adicionales':r,'mu_primaria':mm,'contraste_raw_derivado':primarychecks,
 'asignaturas_plan_concordante':planmatches,'asignaturas_plan_distinto':conflicts,'mismo_plan_observado_ambos_anos_codigo':sameplans,
 'fecha_efectiva_cambio_plan':None,'codigo_nuevo_adoptado':None,'estado':'PENDIENTE_FECHA_EFECTIVA; evidencia anual y asociaciones de asignaturas separadas','clasificacion':'B/C: observación/cruce; E: fecha efectiva de plan y causa'}
 caseout.append(z)
dump(P/'REVISION_16_CASOS.json',caseout)
summary={'personas_cambio_anual':12,'personas_pista_fecha':4,'personas_union':len(people),'solapamiento':0,'fuentes_adicionales_leidas':sum(x['estado']=='LEIDO' for x in inventory),'filas_adicionales_foco':len(rows),
 'filas_primarias_mu':len(mu),'personas_con_asignaturas_plan_confirmado':sum(bool(x['asignaturas_plan_concordante']) for x in caseout),
 'personas_con_plan_discordante_en_asignaturas':sum(bool(x['asignaturas_plan_distinto']) for x in caseout),
 'cambios_codigo_con_mismo_plan_ambos_anos':sum(bool(x['mismo_plan_observado_ambos_anos_codigo']) for x in caseout),
 'controles_raw_derivado':sum(len(x['contraste_raw_derivado']) for x in caseout),'discordancias_raw_derivado':sum(not a['coinciden'] for x in caseout for a in x['contraste_raw_derivado']),
 'eventos_efectivos_cambio_plan_localizados':0,'nuevos_codigos_adoptados':0,'denominadores_conservados':[5901,6622,24774],
 'errores_lectura':len(errors),'clasificacion':'B/C conteos; E fecha efectiva no localizada en fuentes leídas'}
dump(O/'RESULTADOS_PRIMERA_LECTURA.json',summary);dump(O/'INVENTARIO_ADICIONAL.json',inventory);dump(O/'FUENTES_HASHES.json',hashes)
assert all(sha(p)==h for p,h in hashes.items())
print(json.dumps({'resultado':summary,'salida':str(O)},ensure_ascii=False),flush=True)
