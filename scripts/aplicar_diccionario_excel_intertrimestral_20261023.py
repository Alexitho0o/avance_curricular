"""Aplica la consulta canónica a la población MU/T1; no modifica tasas ni fuentes."""
import sys,json,re,collections,datetime,subprocess,os
from pathlib import Path
import openpyxl
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import DiccionarioSiesSeguimiento,rut_normalizado,sha256
stamp=sys.argv[1] if len(sys.argv)>1 else datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
O=R/'outputs'/('aplicacion_diccionario_intertrimestral_'+stamp);O.mkdir(exist_ok=True)
P=Path.home()/'.local/share/ipss'/O.name
assert not P.is_relative_to(R)
assert subprocess.run(['git','-C',str(P.parent),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
P.mkdir(mode=0o700,exist_ok=True);os.chmod(P,0o700)
def dump(p,x,private=False):
 p.write_text(json.dumps(x,ensure_ascii=False,indent=2));
 if private:os.chmod(p,0o600)
source=R/'outputs/reproduccion_puente_20261023_004044/RETENCION_IPSS_PUENTE_REPRODUCIDO_20261023_V11.xlsx'
original=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
mu=R/'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
assert sha256(source)=='dd6d00cf7d329cc1123419206ce8d5f7a1c25b35b452e143153c637ea319772b'
assert sha256(original)=='20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8'
assert sha256(mu)=='19cf33ae4925a7c13ec4745394141a297f4f64be4052a240cb0d2c42ffeb53bd'
git={'rama':subprocess.check_output(['git','branch','--show-current'],cwd=R,text=True).strip(),'head':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip(),'status':subprocess.check_output(['git','status','--short'],cwd=R,text=True)}
dump(O/'GIT_ANTES.json',git)
w=openpyxl.load_workbook(original,read_only=True,data_only=True);v=openpyxl.load_workbook(source,read_only=True,data_only=True)
rows20=[(i,r) for i,r in enumerate(w['20_Foto_BBDD'].iter_rows(min_row=6,values_only=True),6) if r[1]]
a=collections.Counter(tuple(r[j] for j in [1,2,9,10,11]) for i,r in rows20)
b={tuple(r[:5]):r[5] for r in v['20_Foto_BBDD'].iter_rows(min_row=5,values_only=True) if r[0]}
assert a==b
raw9={i:r for i,r in enumerate(w['9_Base_Analitica_Detalle'].iter_rows(values_only=True),1)}
raw11={i:r for i,r in enumerate(w['11_BASE_RETENCION_MU_RAW'].iter_rows(values_only=True),1)}
M=openpyxl.load_workbook(mu,read_only=True,data_only=True)
# Full MU row fingerprints (first 33 columns): book copy is validated, not inferred from the filename.
murows=list(M['BASE_RETENCION_MU'].iter_rows(min_row=2,values_only=True))
muset={tuple('' if x is None else str(x) for x in r[:33]) for r in murows}
mu_typed=collections.defaultdict(set)
for t in murows:
 mu_typed[tuple('' if t[k] is None else str(t[k]) for k in [0,1,2,3])].add('I162S%sC%sJ%sV%s'%tuple(str(t[k]) for k in [11,12,14,15]))
M.close()
s=DiccionarioSiesSeguimiento();records=[json.loads(x[0]) for x in s.db.execute('select payload from observaciones')]
index=collections.defaultdict(list);families=collections.defaultdict(set)
for r in records:
 index[(r['persona'],str(r['ano']))].append(r)
 for e in r['evidencias_matricula']:
  code=e['codigo'];families[code.rsplit('V',1)[0]].add(code)
multi={k for k,x in families.items() if len(x)>1}
photo={}
for year,tab in [(2025,'24_Foto_Corte_T3_2025'),(2026,'25_Foto_Corte_T3_2026')]:
 z=collections.defaultdict(set)
 for r in w[tab].iter_rows(min_row=6,values_only=True):
  if not r[1]:continue
  try:person=rut_normalizado(r[2])
  except ValueError:continue
  z[person].add(str(r[1]))
 photo[year]=z
bridge_manifest=json.loads((R/'outputs/retencion_intermedia_t2_t3/20261021_115156_llaves_20_codcli/MANIFEST_DIAGNOSTICO.json').read_text())
bf=next(x for x in bridge_manifest['fuentes'] if 'IPSS_Procesamiento.xlsx' in x['ruta']);bp=Path(bf['ruta']);assert sha256(bp)==bf['sha256']
bw=openpyxl.load_workbook(bp,read_only=True,data_only=True);bridge=collections.defaultdict(set)
for rr in bw['2025_carga'].iter_rows(min_row=2,values_only=True):
 if len(rr)<74 or rr[72] is None:continue
 try:alias=rut_normalizado(rr[9])
 except ValueError:continue
 bridge[str(rr[72])].add(alias)
bw.close()
origin=[];follow=[];missing=[];seen=set();counts=collections.Counter();other=0;bridged=0
for excelrow,r in rows20:
 if r[2]!=1:continue
 year=int(r[1]);srcrow=int(re.search(r'fila (\d+)',r[3]).group(1));name=str(r[9]);family='';person=None
 try:person=rut_normalizado(r[4])
 except ValueError:pass
 if '9_Base' in r[3]:
  z=raw9[srcrow];code=str(z[23]);typ=str(z[6]);assert z[1]==year
  literal=s.codigos_mu(persona=person,anio=year) if person and typ in ('R','1','RUT','') else {'codigos':[],'evidencias':[]}
  backed=code in mu_typed[tuple('' if z[k] is None else str(z[k]) for k in [1,6,7,8])]
 else:
  z=raw11[srcrow];assert int(z[0])==year
  code='I162S%sC%sJ%sV%s'%tuple(str(z[k]) for k in [11,12,14,15]);typ=str(z[1])
  backed=tuple('' if x is None else str(x) for x in z[:33]) in muset
  literal=s.codigos_mu(persona=person,anio=year) if person and typ in ('R','1','RUT','') else {'codigos':[],'evidencias':[]}
 if typ=='P' and year==2025:
  aliases=bridge.get(str(z[7]),set())
  if len(aliases)==1:person=next(iter(aliases));bridged+=1
 family=code.rsplit('V',1)[0] if re.fullmatch(r'I162S\d+C\d+J\d+V\d+',code) else ''
 assert family and backed, 'Origen MU no respaldado: detener aplicación'
 assert (year,r[8]) not in seen;seen.add((year,r[8]))
 item={'fila_foto':excelrow,'persona':person,'documento_original':r[4],'tipo_doc':typ,'anio':year,'carrera':name,'codigo_origen':code,'familia':family,'multiple':family in multi,'fuente':r[3],'respaldo':'B/C Código literal MU comprobado'}
 matches=[]
 if person:
  for ob in index.get((person,str(year)),[]):
   codes=set(ob['codigos_anuales_documentados'])|set(ob['candidatos_calculados'])
   codes|={str(ob.get(k) or '') for k in ['codigo_documental_atribuible','codigo_interno_adoptado','codigo_propuesto','codigo_propuesto_inicial']}
   fs={c.rsplit('V',1)[0] for c in codes if re.fullmatch(r'I162S\d+C\d+J\d+V\d+',c)}
   relation='Misma familia' if family in fs else 'Otra familia' if fs else 'Familia no determinada'
   if relation!='Misma familia':other+=1
   q=s.consultar(persona=person,codcli=ob['CODCLI'],plan=ob['plan'],anio=ob['ano'],periodo=ob['periodo'],carrera=ob['carrera'],sede=ob['sede_dimension'],jornada=ob['jornada_sies'])
   selected=q['codigo_sies']
   # Do not attach an already resolved OTHER family to the origin career.
   if selected and selected.rsplit('V',1)[0]!=family:relation='Otra familia'
   status='B/C Documental' if q['nivel_respaldo'].startswith('B/C') else 'D Interno aprobado' if q['nivel_respaldo'].startswith('D_') else 'E Pendiente'
   f={**item,'observacion_id':ob['observacion_id'],'codcli':ob['CODCLI'],'plan':ob['plan'],'carrera_umas':ob['carrera'],'periodo_nativo':ob['periodo'],'relacion_familia':relation,'codigo_contexto':selected,'estado':status,'regla':q.get('regla') or q['motivo'],'pendiente':q.get('pendiente') or q['motivo'],'candidatos_no_adoptados':sorted(c for c in codes if re.fullmatch(r'I162S\d+C\d+J\d+V\d+',c)),'codcli_en_foto_t3':('Sí' if ob['CODCLI'] in photo.get(year,{}).get(person,set()) else 'No localizado') if year in photo else 'Sin foto T3 2024','evidencia':q.get('evidencia')}
   matches.append(f);follow.append(f)
 item['contextos']=len(matches);item['contextos_pendientes']=sum(x['estado']=='E Pendiente' for x in matches);origin.append(item)
 if not matches:missing.append(item)
 counts[year]+=1
assert dict(counts)=={2024:558,2025:1150,2026:1471}
assert len({x['observacion_id'] for x in follow})==len(follow)
# Private reproducibility only; no individuals enter workbook or repository JSON.
dump(P/'ORIGENES_T1.json',origin,True);dump(P/'CONSULTAS_EXACTAS.json',follow,True);dump(P/'SIN_CONTEXTO.json',missing,True)
co=collections.Counter((x['anio'],x['carrera'],x['familia'],x['codigo_origen'],'Sí' if x['multiple'] else 'No') for x in origin)
cs=collections.Counter((x['anio'],x['carrera'],x['familia'],x['codigo_origen'],'Sí' if x['multiple'] else 'No') for x in missing)
origin_rows=[list(k)+[n,n-cs[k],cs[k],'B/C'] for k,n in sorted(co.items())]
cf=collections.Counter((x['anio'],x['periodo_nativo'],x['carrera'],x['carrera_umas'],x['plan'],x['familia'],'Sí' if x['multiple'] else 'No',x['codigo_origen'],x['codigo_contexto'] or 'Pendiente',x['estado'],x['codcli_en_foto_t3'],x['relacion_familia'],' | '.join(x['candidatos_no_adoptados']) if x['estado']=='E Pendiente' else '',x['pendiente'] if x['estado']=='E Pendiente' else 'Atribución anual reutilizada; no acredita aceptación SIES ni fecha de cambio') for x in follow)
follow_rows=[list(k[:12])+[n]+list(k[12:]) for k,n in sorted(cf.items())]
st=collections.Counter(x['estado'] for x in follow);fm=collections.Counter(x['estado'] for x in follow if x['multiple'])
summary={'universo_T1_filas_persona_anio':dict(counts),'total_T1':len(origin),'origenes_confirmados':len(origin),'contextos_consultados':len(follow),'estados_contextos':dict(st),'foco_varias_versiones':dict(fm),'personas_anio_sin_contexto':len(missing),'personas_anio_con_contexto':len(origin)-len(missing),'personas_anio_con_pendiente_contexto':sum(x['contextos_pendientes']>0 for x in origin),'foco_origenes_multiples':sum(x['multiple'] for x in origin),'foco_sin_contexto':sum(x['multiple'] for x in missing),'contextos_otra_familia_o_indeterminados_separados':other,'documentos_P_con_puente_explicito':bridged,'estados_por_relacion_familia':dict(collections.Counter(x['relacion_familia']+' / '+x['estado'] for x in follow)),'grupos_liniaje_hoja20':len(a),'filas_liniaje_hoja20':sum(a.values()),'nuevas_adjudicaciones':0}
dump(O/'RESULTADOS.json',summary)
dump(O/'tablas.json',{'origenes':origin_rows,'contextos':follow_rows})
meta={'fecha':datetime.datetime.now().astimezone().isoformat(),'clasificacion':'C aplicación; B observado; D decisiones ya aprobadas; E pendientes','fuentes':{str(p):sha256(p) for p in [source,original,mu,bp,R/'data/diccionario_sies_seguimiento/activo.json',R/'data/diccionario_sies_seguimiento/v1_20261023/descriptor.json']},'diccionario':s.edicion,'origen_excel':str(source),'privado':str(P),'auditoria_privada_hashes':{p.name:sha256(p) for p in P.iterdir()},'grupos_aplicados':'Grupo 1 / MU-T1 2024–2026. Grupos 2 y 3 del libro preservados, no reclasificados.','limites':['Período nativo no equivale automáticamente a corte T2/T3.','Código de origen permanece aunque no se localice contexto.','No localizado en foto T3 no equivale a deserción.','No acredita aceptación SIES.','Las 1.471 filas T1 2026 pertenecen al libro vigente; se conserva ese denominador, distinto de 1.467 citado en una entrega histórica.']}
dump(O/'MANIFEST_APLICACION.json',meta);s.close();w.close();v.close()
print(json.dumps({'output':str(O),'summary':summary},ensure_ascii=False))
