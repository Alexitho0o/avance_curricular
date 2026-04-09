#!/usr/bin/env python3
"""
Postprocesador independiente: BASE_DINAMICA_RESPUESTAS y RESUMEN_RESPUESTAS.

Agrega dos hojas al Excel de caracterización ya generado sin ejecutar el
pipeline principal de caracterización.

Uso:
    python postprocesar_base_dinamica_respuestas.py [ruta_excel]
"""

import sys
import os
import argparse
import glob
import shutil
import subprocess

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

# ─── Configuración ────────────────────────────────────────────────────────────

DEFAULT_INPUT = "/Users/alexi/Desktop/PREGRADO_TRAYECTORIA_CARACTERIZACION_2026_P1_20260512.xlsx"
DESKTOP = os.path.expanduser("~/Desktop")
MAX_EXCEL_ROWS = 1_048_575  # Excel limit minus header

# Columnas que NO pasan a formato largo (identificación + segmentación)
BASE_COLS = frozenset({
    "RUT", "DV", "TOKEN", "RUT_Token", "NOMBRE_COMPLETO", "SEXO",
    "Sector", "SECTOR", "SECTOR_ACADEMICO",
    "Carrera", "CARRERA", "NOMBRE_CARRERA", "COD_CARRERA",
    "Jornada", "JORNADA", "SEDE",
    "ESTADO_ACADEMICO", "MATRICULA", "INSCRITO_LIMESURVEY",
    "RESPONDIO_ENCUESTA", "ESTADO_PARTICIPACION",
    "FUENTE_DATOS", "OBSERVACION_CRUCE",
    # Datos personales: no son preguntas de encuesta
    "CORREO_PERSONAL", "CORREO_INSTITUCIONAL",
})

SECTOR_CANDIDATES  = ["Sector", "SECTOR", "SECTOR_ACADEMICO"]
CARRERA_CANDIDATES = ["Carrera", "CARRERA", "NOMBRE_CARRERA"]
JORNADA_CANDIDATES = ["Jornada", "JORNADA"]
SEXO_CANDIDATES    = ["SEXO", "sexo", "Genero", "GÉNERO"]

# Valores a tratar como vacío (0, 'No', 'N', 'False' son respuestas válidas)
NULL_STRINGS = frozenset({
    "", "nan", "NaN", "None", "none", "<NA>",
    "SIN_DATO", "sin_dato", "N/A", "n/a", "NULL", "null",
})

# ─── Logging ──────────────────────────────────────────────────────────────────

def log(msg):   print(f"[INFO]  {msg}", flush=True)
def warn(msg):  print(f"[WARN]  {msg}", flush=True)
def abort(msg): print(f"[ERROR] {msg}", flush=True); sys.exit(1)

# ─── Utilidades ───────────────────────────────────────────────────────────────

def find_col(columns, candidates):
    col_set = set(columns)
    for c in candidates:
        if c in col_set:
            return c
    return None


def discover_file(path):
    """Devuelve la ruta real del archivo; lo busca por nombre si no existe."""
    if os.path.exists(path):
        return path
    fname = os.path.basename(path)
    search_roots = [
        DESKTOP,
        os.path.expanduser("~/Downloads"),
        os.path.expanduser("~/Documents"),
        os.path.expanduser("~/Library/CloudStorage"),
    ]
    for root in search_roots:
        if not os.path.isdir(root):
            continue
        found = glob.glob(os.path.join(root, "**", fname), recursive=True)
        if found:
            found.sort(key=os.path.getmtime, reverse=True)
            return found[0]
    return None


def find_sector_file():
    """Localiza el archivo ANALISIS_SECTORIZADO en ubicaciones comunes."""
    patterns = [
        os.path.expanduser("~/Library/CloudStorage/**/*ANALISIS_SECTORIZADO*.xlsx"),
        os.path.expanduser("~/Desktop/*ANALISIS_SECTORIZADO*.xlsx"),
        os.path.expanduser("~/Downloads/*ANALISIS_SECTORIZADO*.xlsx"),
        os.path.expanduser("~/Documents/**/*ANALISIS_SECTORIZADO*.xlsx"),
    ]
    found = []
    for p in patterns:
        found.extend(glob.glob(p, recursive=True))
    if found:
        found.sort(key=os.path.getmtime, reverse=True)
        return found[0]
    return None


def compute_col_widths(df, headers=None):
    """Calcula anchos de columna desde pandas (vectorizado, rápido)."""
    widths = {}
    cols = headers or df.columns
    for col in cols:
        header_len = len(str(col))
        if col in df.columns:
            try:
                p95 = df[col].dropna().astype(str).str.len().quantile(0.95)
                max_len = max(header_len, int(p95) if not pd.isna(p95) else header_len)
            except Exception:
                max_len = header_len
        else:
            max_len = header_len
        widths[col] = max(10, min(max_len + 2, 60))
    return widths


def format_sheet(ws, col_widths, table_name=None):
    """Aplica encabezado coloreado, freeze, autofilter y anchos de columna."""
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col_idx, (col_name, width) in enumerate(col_widths.items(), 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if table_name:
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row > 1 and max_col > 0:
            try:
                tbl_ref = f"A1:{get_column_letter(max_col)}{max_row}"
                tbl = Table(displayName=table_name, ref=tbl_ref)
                tbl.tableStyleInfo = TableStyleInfo(
                    name="TableStyleMedium9",
                    showRowStripes=True,
                    showColumnStripes=False,
                    showFirstColumn=False,
                    showLastColumn=False,
                )
                ws.add_table(tbl)
            except Exception as e:
                warn(f"No se pudo crear tabla estructurada '{table_name}': {e}")

# ─── Enriquecimiento de Sector ────────────────────────────────────────────────

def enriquecer_sector(df, carrera_col):
    """
    Busca la columna Sector en el DataFrame.
    Si no existe, intenta unirla desde el archivo ANALISIS_SECTORIZADO.
    Retorna (df_enriquecido, nombre_columna_sector, fuente).
    """
    col = find_col(df.columns, SECTOR_CANDIDATES)
    if col:
        return df, col, "CONSOLIDADO_RUT"

    warn("Columna de sector no encontrada en CONSOLIDADO_RUT. Buscando en archivo externo…")
    sec_file = find_sector_file()

    if not sec_file:
        warn("No se encontró archivo ANALISIS_SECTORIZADO. Sector quedará vacío.")
        df["Sector"] = None
        return df, "Sector", "ninguna (no encontrado)"

    log(f"Archivo de sector: {sec_file}")

    # Intento 1: BASE_SECTORIZADA  (RUT -> SECTOR_ACADEMICO)
    # CONSOLIDADO_RUT tiene RUT solo numérico (ej. '16560868').
    # BASE_SECTORIZADA tiene RUT con guion y DV (ej. '16560868-2').
    # Se normaliza extrayendo la parte antes del guion.
    try:
        df_sec = pd.read_excel(sec_file, sheet_name="BASE_SECTORIZADA", dtype=str)
        if "RUT" in df_sec.columns and "SECTOR_ACADEMICO" in df_sec.columns:
            df_sec = df_sec[["RUT", "SECTOR_ACADEMICO"]].dropna(subset=["RUT"]).copy()
            df_sec["_rut_num"] = df_sec["RUT"].str.split("-").str[0].str.strip()
            df_sec = (df_sec[["_rut_num", "SECTOR_ACADEMICO"]]
                      .drop_duplicates("_rut_num")
                      .rename(columns={"SECTOR_ACADEMICO": "Sector", "_rut_num": "RUT"}))
            df["_rut_num"] = df["RUT"].str.strip()
            df = df.merge(df_sec, left_on="_rut_num", right_on="RUT", how="left", suffixes=("", "_sec"))
            df = df.drop(columns=["_rut_num", "RUT_sec"], errors="ignore")
            n = df["Sector"].notna().sum()
            log(f"Sector unido por RUT (normalizado): {n:,}/{len(df):,} filas con sector asignado")
            return df, "Sector", f"BASE_SECTORIZADA ({os.path.basename(sec_file)})"
    except Exception as e:
        warn(f"No se pudo leer BASE_SECTORIZADA: {e}")

    # Intento 2: SECTOR_CARRERA  (CARRERA -> SECTOR_ACADEMICO)
    if carrera_col:
        try:
            df_sc = pd.read_excel(sec_file, sheet_name="SECTOR_CARRERA", dtype=str)
            if "SECTOR_ACADEMICO" in df_sc.columns and "CARRERA" in df_sc.columns:
                df_sc = (df_sc[["CARRERA", "SECTOR_ACADEMICO"]]
                         .dropna(subset=["CARRERA"])
                         .drop_duplicates("CARRERA"))
                df_sc["_key"] = df_sc["CARRERA"].str.strip().str.upper()
                df["_key"] = df[carrera_col].str.strip().str.upper()
                df = df.merge(
                    df_sc[["_key", "SECTOR_ACADEMICO"]].rename(columns={"SECTOR_ACADEMICO": "Sector"}),
                    on="_key", how="left"
                ).drop(columns=["_key"])
                n = df["Sector"].notna().sum()
                log(f"Sector unido por CARRERA: {n:,}/{len(df):,} filas con sector asignado")
                return df, "Sector", f"SECTOR_CARRERA ({os.path.basename(sec_file)})"
        except Exception as e:
            warn(f"No se pudo leer SECTOR_CARRERA: {e}")

    warn("No se pudo obtener sector desde archivo externo. Sector quedará vacío.")
    df["Sector"] = None
    return df, "Sector", "ninguna (fallo de lectura)"

# ─── Main ─────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Agrega BASE_DINAMICA_RESPUESTAS y RESUMEN_RESPUESTAS al Excel de caracterización."
    )
    parser.add_argument(
        "ruta_excel", nargs="?", default=DEFAULT_INPUT,
        help="Ruta al archivo Excel de caracterización"
    )
    args = parser.parse_args()

    # ── 1. Localizar archivo ──────────────────────────────────────────────────
    input_path = discover_file(args.ruta_excel)
    if not input_path:
        abort(f"No se encontró el archivo: {args.ruta_excel}")
    log(f"Archivo de origen: {input_path}")

    # ── 2. Validar hoja CONSOLIDADO_RUT ──────────────────────────────────────
    wb_check = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    if "CONSOLIDADO_RUT" not in wb_check.sheetnames:
        wb_check.close()
        abort("La hoja CONSOLIDADO_RUT no existe en el archivo.")
    wb_check.close()
    log("Hoja CONSOLIDADO_RUT: encontrada ✓")

    # ── 3. Leer CONSOLIDADO_RUT ───────────────────────────────────────────────
    log("Leyendo CONSOLIDADO_RUT con pandas…")
    df = pd.read_excel(input_path, sheet_name="CONSOLIDADO_RUT", dtype=str)
    n_filas = len(df)
    n_cols_total = len(df.columns)
    log(f"Filas leídas: {n_filas:,}  |  Columnas totales: {n_cols_total}")

    if n_filas == 0:
        abort("CONSOLIDADO_RUT está vacía.")

    # ── 4. Validar columnas clave ─────────────────────────────────────────────
    rut_col = find_col(df.columns, ["RUT"])
    token_col = find_col(df.columns, ["TOKEN", "RUT_Token"])
    if not rut_col and not token_col:
        abort("No se encontró columna RUT ni TOKEN en CONSOLIDADO_RUT.")

    sexo_col   = find_col(df.columns, SEXO_CANDIDATES)
    carrera_col = find_col(df.columns, CARRERA_CANDIDATES)
    jornada_col = find_col(df.columns, JORNADA_CANDIDATES)

    if not sexo_col:    warn("No se encontró columna SEXO.")
    if not carrera_col: warn("No se encontró columna de carrera.")
    if not jornada_col: warn("No se encontró columna de jornada.")

    # ── 5. Enriquecer con Sector ──────────────────────────────────────────────
    df, sector_col, sector_source = enriquecer_sector(df, carrera_col)
    log(f"Columna de sector usada: '{sector_col}' — fuente: {sector_source}")

    # ── 6. Renombrar a nombres canónicos para el output ───────────────────────
    rename_map = {}
    if sector_col  and sector_col  != "Sector":  rename_map[sector_col]  = "Sector"
    if carrera_col and carrera_col != "Carrera": rename_map[carrera_col] = "Carrera"
    if jornada_col and jornada_col != "Jornada": rename_map[jornada_col] = "Jornada"
    if rename_map:
        df = df.rename(columns=rename_map)
    sector_col  = "Sector"
    carrera_col = "Carrera"  if carrera_col else None
    jornada_col = "Jornada"  if jornada_col else None

    # ── 7. Detectar columnas de encuesta ──────────────────────────────────────
    effective_base = BASE_COLS | {"Sector", "Carrera", "Jornada"}
    id_cols     = [c for c in df.columns if c in effective_base]
    survey_cols = [c for c in df.columns if c not in effective_base]

    if not survey_cols:
        abort("No se detectaron columnas de encuesta.")

    log(f"Columnas base retenidas:  {len(id_cols)}")
    log(f"Columnas de encuesta:     {len(survey_cols)}")

    # ── 8. Melt a formato largo ───────────────────────────────────────────────
    log("Transformando a formato largo (melt)… esto puede tardar un momento.")
    df_long = df[id_cols + survey_cols].melt(
        id_vars=id_cols,
        value_vars=survey_cols,
        var_name="Pregunta",
        value_name="Respuesta",
    )

    n_antes = len(df_long)
    null_mask = df_long["Respuesta"].isna() | df_long["Respuesta"].isin(NULL_STRINGS)
    df_long = df_long[~null_mask].copy()
    n_despues = len(df_long)
    log(f"Filas antes del filtro: {n_antes:,}  →  después: {n_despues:,}")

    if n_despues == 0:
        abort("BASE_DINAMICA_RESPUESTAS no tiene filas después del filtro de nulos.")

    if n_despues > MAX_EXCEL_ROWS:
        warn(f"La base larga tiene {n_despues:,} filas. Excel solo admite {MAX_EXCEL_ROWS:,}.")
        warn("Se truncará. Considere filtrar por RESPONDIO_ENCUESTA='SI' si necesita todas las filas.")
        df_long = df_long.head(MAX_EXCEL_ROWS)

    # ── 9. Ordenar columnas de BASE_DINAMICA_RESPUESTAS ───────────────────────
    priority  = ["RUT", "TOKEN", "SEXO", "Sector", "Carrera", "Jornada"]
    optional  = ["COD_CARRERA", "SEDE", "ESTADO_ACADEMICO", "RESPONDIO_ENCUESTA", "ESTADO_PARTICIPACION"]
    tail      = ["Pregunta", "Respuesta"]
    col_order = [c for c in priority + optional + tail if c in df_long.columns]
    df_base   = df_long[col_order].reset_index(drop=True)

    preguntas_unicas  = df_base["Pregunta"].nunique()
    respuestas_unicas = df_base["Respuesta"].nunique()
    log(f"Preguntas únicas:  {preguntas_unicas:,}")
    log(f"Respuestas únicas: {respuestas_unicas:,}")

    # ── 10. Crear RESUMEN_RESPUESTAS ──────────────────────────────────────────
    log("Creando RESUMEN_RESPUESTAS…")
    group_cols = [c for c in ["Pregunta", "Respuesta", "Sector", "Carrera", "Jornada", "SEXO"]
                  if c in df_base.columns]
    count_col = "RUT" if "RUT" in df_base.columns else "TOKEN"
    df_resumen = (
        df_base.groupby(group_cols, dropna=False)[count_col]
        .nunique()
        .reset_index(name="TOTAL_RUT")
    )
    log(f"Filas en RESUMEN_RESPUESTAS: {len(df_resumen):,}")

    if len(df_resumen) == 0:
        abort("RESUMEN_RESPUESTAS no tiene filas.")

    # ── 11. Preparar ruta de salida ───────────────────────────────────────────
    stem        = os.path.splitext(os.path.basename(input_path))[0]
    output_name = f"{stem}_CON_BASE_DINAMICA.xlsx"
    output_path = os.path.join(DESKTOP, output_name)
    log(f"Archivo de salida: {output_path}")

    shutil.copy2(input_path, output_path)
    log("Copia del archivo original creada. ✓")

    # ── 12. Escribir hojas nuevas (pandas ExcelWriter, modo append) ───────────
    log("Escribiendo BASE_DINAMICA_RESPUESTAS y RESUMEN_RESPUESTAS…")
    with pd.ExcelWriter(
        output_path,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        df_base.to_excel(writer, sheet_name="BASE_DINAMICA_RESPUESTAS", index=False)
        df_resumen.to_excel(writer, sheet_name="RESUMEN_RESPUESTAS", index=False)
    log("Datos escritos. ✓")

    # ── 13. Aplicar formato a las hojas nuevas ────────────────────────────────
    log("Aplicando formato (encabezado, freeze, autofilter, anchos)…")
    wb = openpyxl.load_workbook(output_path)

    widths_base    = compute_col_widths(df_base)
    widths_resumen = compute_col_widths(df_resumen)

    format_sheet(wb["BASE_DINAMICA_RESPUESTAS"],    widths_base,    table_name="TBL_BASE_DINAMICA")
    format_sheet(wb["RESUMEN_RESPUESTAS"],          widths_resumen, table_name="TBL_RESUMEN_RESPUESTAS")

    wb.save(output_path)
    wb.close()
    log("Formato aplicado y archivo guardado. ✓")

    # ── 14. Abrir archivo ─────────────────────────────────────────────────────
    try:
        subprocess.Popen(["open", output_path])
        apertura_ok = True
        log("Apertura automática: ✓")
    except Exception as e:
        apertura_ok = False
        warn(f"No se pudo abrir automáticamente: {e}")

    # ── 15. Resumen en log ────────────────────────────────────────────────────
    separador = "=" * 64
    print(f"\n{separador}")
    print("RESUMEN DEL POSTPROCESO")
    print(separador)
    print(f"  Archivo de origen            : {input_path}")
    print(f"  Archivo de salida            : {output_path}")
    print(f"  Filas en CONSOLIDADO_RUT     : {n_filas:,}")
    print(f"  Columnas totales             : {n_cols_total}")
    print(f"  Columna de sector            : {sector_col} ({sector_source})")
    print(f"  Columna de carrera           : {carrera_col or 'no encontrada'}")
    print(f"  Columna de jornada           : {jornada_col or 'no encontrada'}")
    print(f"  Columnas de encuesta         : {len(survey_cols)}")
    print(f"  Filas BASE_DINAMICA          : {len(df_base):,}")
    print(f"  Preguntas únicas             : {preguntas_unicas:,}")
    print(f"  Respuestas únicas            : {respuestas_unicas:,}")
    print(f"  Filas RESUMEN_RESPUESTAS     : {len(df_resumen):,}")
    print(f"  Apertura automática          : {'✓' if apertura_ok else 'falló (abrir manualmente)'}")
    print(separador)


if __name__ == "__main__":
    main()
