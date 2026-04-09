#!/usr/bin/env python3

from pathlib import Path
from datetime import datetime
import hashlib
import json

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


# ============================================================
# CONTEXTO
# ============================================================

RAIZ = Path(
    "/Users/alexi/Documents/GitHub/avance_curricular"
).resolve()

EXPEDIENTE_644 = (
    RAIZ
    / "avance_curricular_2026"
    / "04_gobernanza_mallas"
    / "03_auditorias"
    / "EXPEDIENTE_REVISION_644_20260702_012504"
    / "02_RESULTADOS"
    / "02_REVISION_INDIVIDUAL_PRIORIDAD_1.tsv"
)

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

SALIDA = (
    RAIZ
    / "avance_curricular_2026"
    / "04_gobernanza_mallas"
    / "03_auditorias"
    / f"PLANILLA_DECISION_29_P1_{timestamp}"
)

RESULTADOS = SALIDA / "02_RESULTADOS"
AUDITORIAS = SALIDA / "03_AUDITORIAS"
REPORTES = SALIDA / "04_REPORTES"

for carpeta in (
    RESULTADOS,
    AUDITORIAS,
    REPORTES,
):
    carpeta.mkdir(parents=True, exist_ok=True)


# ============================================================
# UTILIDADES
# ============================================================

def sha256(ruta):
    h = hashlib.sha256()

    with ruta.open("rb") as archivo:
        for bloque in iter(
            lambda: archivo.read(1024 * 1024),
            b"",
        ):
            h.update(bloque)

    return h.hexdigest()


def asegurar_columna(df, columna, valor=""):
    if columna not in df.columns:
        df[columna] = valor


def seleccionar_primera_disponible(df, opciones):
    for columna in opciones:
        if columna in df.columns:
            return columna

    return None


def serie_o_vacio(df, columna):
    if columna and columna in df.columns:
        return (
            df[columna]
            .fillna("")
            .astype(str)
            .str.strip()
        )

    return pd.Series(
        "",
        index=df.index,
        dtype="string",
    )


# ============================================================
# 1. VALIDAR FUENTE
# ============================================================

if not EXPEDIENTE_644.exists():
    raise RuntimeError(
        f"No existe el expediente P1: {EXPEDIENTE_644}"
    )


# ============================================================
# 2. LEER LOS 29 CASOS
# ============================================================

base = pd.read_csv(
    EXPEDIENTE_644,
    sep="\t",
    dtype=str,
    keep_default_na=False,
)

if len(base) != 29:
    raise RuntimeError(
        f"Se esperaban 29 casos P1 y se encontraron "
        f"{len(base)}."
    )

if "ID_FILA_5809" not in base.columns:
    raise RuntimeError(
        "Falta ID_FILA_5809."
    )

if base["ID_FILA_5809"].duplicated().any():
    raise RuntimeError(
        "Existen ID_FILA_5809 duplicados."
    )

prioridades_validas = {
    "P1_CRITICA",
    "P1_ALTA",
}

if not set(
    base["PRIORIDAD_REVISION"].dropna().unique()
).issubset(prioridades_validas):
    raise RuntimeError(
        "El archivo contiene casos fuera de P1."
    )


# ============================================================
# 3. IDENTIFICAR COLUMNAS DE PRESENTACIÓN
# ============================================================

col_documento = seleccionar_primera_disponible(
    base,
    [
        "DOCUMENTO_NORM",
        "NUM_DOCUMENTO",
        "RUT",
    ],
)

col_nombres = seleccionar_primera_disponible(
    base,
    [
        "NOMBRES",
        "NOMBRE",
    ],
)

col_paterno = seleccionar_primera_disponible(
    base,
    [
        "PRIMER_APELLIDO",
        "APELLIDO_PATERNO",
    ],
)

col_materno = seleccionar_primera_disponible(
    base,
    [
        "SEGUNDO_APELLIDO",
        "APELLIDO_MATERNO",
    ],
)

col_codigo_unico = seleccionar_primera_disponible(
    base,
    [
        "CODIGO_UNICO",
        "CODIGO_UNICO_CARRERA",
    ],
)

col_plan = seleccionar_primera_disponible(
    base,
    [
        "CODPESTUD_RESUELTO",
        "CODPESTUD",
    ],
)

col_carrera_plan = seleccionar_primera_disponible(
    base,
    [
        "CODCARR_PLAN_NORM",
        "CODCARR_PLAN",
    ],
)

col_carrera_evidencia = seleccionar_primera_disponible(
    base,
    [
        "CODCARPR_DATOS_NORM",
        "CARRERAS_DATOS_IDENTIDAD",
    ],
)

col_nivel_evidencia = seleccionar_primera_disponible(
    base,
    [
        "NIVEL_CANDIDATO_NUM",
        "NIVEL_CANDIDATO_IDENTIDAD",
        "NIVELES_DATOS_IDENTIDAD",
    ],
)

col_nivel_max_plan = seleccionar_primera_disponible(
    base,
    [
        "NIVEL_MAX_PLAN_NUM",
        "NIVEL_MAX_PLAN_DIAG",
        "NIVEL_MAX_PLAN",
    ],
)

col_planes_historicos = seleccionar_primera_disponible(
    base,
    [
        "PLANES_HISTORICOS",
        "PLANES_DOCUMENTO",
    ],
)

col_carreras_identidad = seleccionar_primera_disponible(
    base,
    [
        "CARRERAS_DATOS_IDENTIDAD",
        "CARRERAS_DATOS",
    ],
)

col_niveles_identidad = seleccionar_primera_disponible(
    base,
    [
        "NIVELES_DATOS_IDENTIDAD",
        "NIVELES_DATOS",
    ],
)

col_rut_identidad = seleccionar_primera_disponible(
    base,
    [
        "RUT_DATOS_IDENTIDAD",
        "DOCUMENTOS_DATOS_CODCLI",
    ],
)

col_codcli_identidad = seleccionar_primera_disponible(
    base,
    [
        "CODCLI_DATOS_IDENTIDAD",
        "CODCLI_DATOS_DOCUMENTO",
    ],
)


# ============================================================
# 4. CONSTRUIR PLANILLA OPERATIVA
# ============================================================

planilla = pd.DataFrame({
    "N_ORDEN_REVISION":
    pd.to_numeric(
        base.get(
            "N_ORDEN_REVISION",
            pd.Series(
                range(1, len(base) + 1),
                index=base.index,
            ),
        ),
        errors="coerce",
    ).fillna(0).astype(int),

    "ID_FILA_5809":
    serie_o_vacio(base, "ID_FILA_5809"),

    "PRIORIDAD_REVISION":
    serie_o_vacio(base, "PRIORIDAD_REVISION"),

    "CATEGORIA_REVISION":
    serie_o_vacio(base, "CATEGORIA_REVISION_644"),

    "DOCUMENTO":
    serie_o_vacio(base, col_documento),

    "NOMBRES":
    serie_o_vacio(base, col_nombres),

    "PRIMER_APELLIDO":
    serie_o_vacio(base, col_paterno),

    "SEGUNDO_APELLIDO":
    serie_o_vacio(base, col_materno),

    "CODIGO_UNICO_5809":
    serie_o_vacio(base, col_codigo_unico),

    "CODPESTUD_VIGENTE":
    serie_o_vacio(base, col_plan),

    "CODCARR_PLAN_VIGENTE":
    serie_o_vacio(base, col_carrera_plan),

    "CODCARPR_EVIDENCIA":
    serie_o_vacio(base, col_carrera_evidencia),

    "NIVEL_EVIDENCIA":
    serie_o_vacio(base, col_nivel_evidencia),

    "NIVEL_MAX_PLAN":
    serie_o_vacio(base, col_nivel_max_plan),

    "PLANES_HISTORICOS":
    serie_o_vacio(base, col_planes_historicos),

    "RUT_EVIDENCIA_IDENTIDAD":
    serie_o_vacio(base, col_rut_identidad),

    "CODCLI_EVIDENCIA_IDENTIDAD":
    serie_o_vacio(base, col_codcli_identidad),

    "CARRERAS_EVIDENCIA_IDENTIDAD":
    serie_o_vacio(base, col_carreras_identidad),

    "NIVELES_EVIDENCIA_IDENTIDAD":
    serie_o_vacio(base, col_niveles_identidad),

    "RESPALDO_ACTUAL":
    serie_o_vacio(base, "RESPALDO_ACTUAL"),

    "OBSERVACION_DIAGNOSTICO":
    serie_o_vacio(base, "OBSERVACION_EXPEDIENTE"),

    "ACCION_REQUERIDA":
    serie_o_vacio(base, "ACCION_REQUERIDA"),

    "DECISION_MANUAL":
    "",

    "NIVEL_APROBADO":
    "",

    "PLAN_APROBADO":
    "",

    "CARRERA_APROBADA":
    "",

    "FUENTE_RESPALDO_DECISION":
    "",

    "FUNDAMENTO_DECISION":
    "",

    "RESPONSABLE_REVISION":
    "",

    "FECHA_REVISION":
    "",

    "ESTADO_CIERRE":
    "PENDIENTE_REVISION",

    "OBSERVACIONES_REVISOR":
    "",
})


# ============================================================
# 5. VALIDAR COMPOSICIÓN
# ============================================================

conteos = (
    planilla["CATEGORIA_REVISION"]
    .value_counts()
    .to_dict()
)

esperados = {
    "CONFLICTO_CARRERA_DATOSALUMNOS_VS_PLAN": 14,
    "IDENTIDAD_EXACTA_MULTIPLES_TRAYECTORIAS": 14,
    "NIVEL_OBSERVADO_SUPERA_MAXIMO_PLAN": 1,
}

diferencias = []

for categoria, esperado in esperados.items():
    observado = int(
        conteos.get(categoria, 0)
    )

    if observado != esperado:
        diferencias.append({
            "CATEGORIA": categoria,
            "ESPERADO": esperado,
            "OBSERVADO": observado,
        })

if diferencias:
    pd.DataFrame(diferencias).to_csv(
        AUDITORIAS
        / "00_DIFERENCIAS_CATEGORIAS.tsv",
        sep="\t",
        index=False,
    )

    raise RuntimeError(
        "La composición de los 29 P1 no coincide con "
        "la distribución confirmada."
    )


# ============================================================
# 6. GUÍA DE DECISIÓN
# ============================================================

guia = pd.DataFrame([
    {
        "CAMPO": "DECISION_MANUAL",
        "VALORES_PERMITIDOS": (
            "APROBAR_NIVEL | RECHAZAR_NIVEL | "
            "MANTENER_PENDIENTE"
        ),
        "REGLA_OPERATIVA": (
            "APROBAR_NIVEL solo cuando exista evidencia "
            "institucional suficiente del plan, carrera y nivel."
        ),
    },
    {
        "CAMPO": "NIVEL_APROBADO",
        "VALORES_PERMITIDOS": "Número entero mayor o igual a 1",
        "REGLA_OPERATIVA": (
            "Obligatorio únicamente cuando DECISION_MANUAL "
            "sea APROBAR_NIVEL."
        ),
    },
    {
        "CAMPO": "PLAN_APROBADO",
        "VALORES_PERMITIDOS": "CODPESTUD confirmado",
        "REGLA_OPERATIVA": (
            "Debe coincidir con el plan institucional respaldado."
        ),
    },
    {
        "CAMPO": "CARRERA_APROBADA",
        "VALORES_PERMITIDOS": "Código de carrera confirmado",
        "REGLA_OPERATIVA": (
            "Debe corresponder al plan aprobado."
        ),
    },
    {
        "CAMPO": "FUENTE_RESPALDO_DECISION",
        "VALORES_PERMITIDOS": (
            "Nombre, ruta, hoja, reporte o evidencia institucional"
        ),
        "REGLA_OPERATIVA": (
            "No puede quedar vacío cuando exista una decisión final."
        ),
    },
    {
        "CAMPO": "FUNDAMENTO_DECISION",
        "VALORES_PERMITIDOS": "Texto libre verificable",
        "REGLA_OPERATIVA": (
            "Debe explicar por qué la evidencia permite aprobar, "
            "rechazar o mantener pendiente."
        ),
    },
    {
        "CAMPO": "ESTADO_CIERRE",
        "VALORES_PERMITIDOS": (
            "PENDIENTE_REVISION | EN_REVISION | "
            "CERRADO_APROBADO | CERRADO_RECHAZADO | "
            "CERRADO_SIN_EVIDENCIA"
        ),
        "REGLA_OPERATIVA": (
            "Solo usar CERRADO_APROBADO cuando todos los campos "
            "obligatorios estén completos."
        ),
    },
])


# ============================================================
# 7. RESÚMENES
# ============================================================

resumen_categorias = (
    planilla.groupby(
        [
            "PRIORIDAD_REVISION",
            "CATEGORIA_REVISION",
            "ACCION_REQUERIDA",
        ],
        dropna=False,
    )
    .size()
    .reset_index(name="CASOS")
)

resumen_kpi = pd.DataFrame([
    {
        "INDICADOR": "Casos P1",
        "VALOR": 29,
    },
    {
        "INDICADOR": "Prioridad crítica",
        "VALOR": int(
            planilla[
                "PRIORIDAD_REVISION"
            ].eq("P1_CRITICA").sum()
        ),
    },
    {
        "INDICADOR": "Prioridad alta",
        "VALOR": int(
            planilla[
                "PRIORIDAD_REVISION"
            ].eq("P1_ALTA").sum()
        ),
    },
    {
        "INDICADOR": "Decisiones registradas inicialmente",
        "VALOR": 0,
    },
    {
        "INDICADOR": "Casos aprobados inicialmente",
        "VALOR": 0,
    },
    {
        "INDICADOR": "Fuente original modificada",
        "VALOR": "NO",
    },
    {
        "INDICADOR": "Archivo final de carga generado",
        "VALOR": "NO",
    },
])


# ============================================================
# 8. EXPORTAR TSV
# ============================================================

planilla.to_csv(
    RESULTADOS
    / "01_PLANILLA_OPERATIVA_DECISION_29_P1.tsv",
    sep="\t",
    index=False,
)

guia.to_csv(
    RESULTADOS
    / "02_GUIA_DECISION_29_P1.tsv",
    sep="\t",
    index=False,
)

resumen_categorias.to_csv(
    AUDITORIAS
    / "01_RESUMEN_CATEGORIAS.tsv",
    sep="\t",
    index=False,
)


# ============================================================
# 9. VALIDACIONES
# ============================================================

validaciones = pd.DataFrame([
    {
        "VALIDACION": "UNIVERSO_29",
        "RESULTADO": (
            "OK" if len(planilla) == 29 else "ERROR"
        ),
        "DETALLE": len(planilla),
    },
    {
        "VALIDACION": "ID_FILA_UNICO",
        "RESULTADO": (
            "OK"
            if not planilla[
                "ID_FILA_5809"
            ].duplicated().any()
            else "ERROR"
        ),
        "DETALLE": int(
            planilla[
                "ID_FILA_5809"
            ].duplicated().sum()
        ),
    },
    {
        "VALIDACION": "COMPOSICION_ESPERADA",
        "RESULTADO": "OK",
        "DETALLE": "14 + 14 + 1",
    },
    {
        "VALIDACION": "DECISIONES_AUTOMATICAS",
        "RESULTADO": "OK",
        "DETALLE": 0,
    },
    {
        "VALIDACION": "ESTADO_INICIAL",
        "RESULTADO": (
            "OK"
            if planilla[
                "ESTADO_CIERRE"
            ].eq(
                "PENDIENTE_REVISION"
            ).all()
            else "ERROR"
        ),
        "DETALLE": int(
            planilla[
                "ESTADO_CIERRE"
            ].eq(
                "PENDIENTE_REVISION"
            ).sum()
        ),
    },
    {
        "VALIDACION": "FUENTES_MODIFICADAS",
        "RESULTADO": "OK",
        "DETALLE": "NO",
    },
    {
        "VALIDACION": "ARCHIVO_CARGA_GENERADO",
        "RESULTADO": "OK",
        "DETALLE": "NO",
    },
])

validaciones.to_csv(
    AUDITORIAS
    / "02_VALIDACIONES.tsv",
    sep="\t",
    index=False,
)


# ============================================================
# 10. FUENTE Y HASH
# ============================================================

fuentes_df = pd.DataFrame([
    {
        "ROL": "EXPEDIENTE_P1",
        "RUTA": str(EXPEDIENTE_644),
        "SHA256": sha256(EXPEDIENTE_644),
        "TAMANO_BYTES": EXPEDIENTE_644.stat().st_size,
    }
])

fuentes_df.to_csv(
    AUDITORIAS
    / "03_FUENTE_Y_HASH.tsv",
    sep="\t",
    index=False,
)


# ============================================================
# 11. EXCEL OPERATIVO
# ============================================================

excel = (
    RESULTADOS
    / "PLANILLA_OPERATIVA_DECISION_29_P1.xlsx"
)

with pd.ExcelWriter(
    excel,
    engine="openpyxl",
) as writer:
    resumen_kpi.to_excel(
        writer,
        sheet_name="RESUMEN",
        index=False,
    )

    guia.to_excel(
        writer,
        sheet_name="GUIA_DECISION",
        index=False,
    )

    planilla.to_excel(
        writer,
        sheet_name="REVISION_29",
        index=False,
    )

    planilla[
        planilla[
            "CATEGORIA_REVISION"
        ].eq(
            "NIVEL_OBSERVADO_SUPERA_MAXIMO_PLAN"
        )
    ].to_excel(
        writer,
        sheet_name="CASO_CRITICO",
        index=False,
    )

    planilla[
        planilla[
            "CATEGORIA_REVISION"
        ].eq(
            "CONFLICTO_CARRERA_DATOSALUMNOS_VS_PLAN"
        )
    ].to_excel(
        writer,
        sheet_name="CONFLICTO_CARRERA",
        index=False,
    )

    planilla[
        planilla[
            "CATEGORIA_REVISION"
        ].eq(
            "IDENTIDAD_EXACTA_MULTIPLES_TRAYECTORIAS"
        )
    ].to_excel(
        writer,
        sheet_name="MULTIPLES_TRAYECTORIAS",
        index=False,
    )

    validaciones.to_excel(
        writer,
        sheet_name="VALIDACIONES",
        index=False,
    )

    fuentes_df.to_excel(
        writer,
        sheet_name="FUENTES",
        index=False,
    )

    libro = writer.book

    # Listas auxiliares ocultas para validaciones.
    listas = libro.create_sheet("_LISTAS")

    decisiones = [
        "APROBAR_NIVEL",
        "RECHAZAR_NIVEL",
        "MANTENER_PENDIENTE",
    ]

    estados = [
        "PENDIENTE_REVISION",
        "EN_REVISION",
        "CERRADO_APROBADO",
        "CERRADO_RECHAZADO",
        "CERRADO_SIN_EVIDENCIA",
    ]

    for i, valor in enumerate(
        decisiones,
        start=1,
    ):
        listas.cell(
            row=i,
            column=1,
            value=valor,
        )

    for i, valor in enumerate(
        estados,
        start=1,
    ):
        listas.cell(
            row=i,
            column=2,
            value=valor,
        )

    listas.sheet_state = "hidden"

    hoja_revision = libro["REVISION_29"]

    encabezados = {
        celda.value: celda.column
        for celda in hoja_revision[1]
    }

    col_decision = encabezados["DECISION_MANUAL"]
    col_nivel = encabezados["NIVEL_APROBADO"]
    col_estado = encabezados["ESTADO_CIERRE"]

    letra_decision = hoja_revision.cell(
        row=1,
        column=col_decision,
    ).column_letter

    letra_nivel = hoja_revision.cell(
        row=1,
        column=col_nivel,
    ).column_letter

    letra_estado = hoja_revision.cell(
        row=1,
        column=col_estado,
    ).column_letter

    dv_decision = DataValidation(
        type="list",
        formula1="'_LISTAS'!$A$1:$A$3",
        allow_blank=True,
    )

    dv_estado = DataValidation(
        type="list",
        formula1="'_LISTAS'!$B$1:$B$5",
        allow_blank=False,
    )

    dv_nivel = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="99",
        allow_blank=True,
    )

    hoja_revision.add_data_validation(
        dv_decision
    )

    hoja_revision.add_data_validation(
        dv_estado
    )

    hoja_revision.add_data_validation(
        dv_nivel
    )

    dv_decision.add(
        f"{letra_decision}2:"
        f"{letra_decision}{len(planilla) + 1}"
    )

    dv_estado.add(
        f"{letra_estado}2:"
        f"{letra_estado}{len(planilla) + 1}"
    )

    dv_nivel.add(
        f"{letra_nivel}2:"
        f"{letra_nivel}{len(planilla) + 1}"
    )

    # Formato general.
    fill_encabezado = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    fill_decision = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    fill_critico = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )

    fill_alta = PatternFill(
        fill_type="solid",
        fgColor="FCE5CD",
    )

    for hoja in libro.worksheets:
        if hoja.title == "_LISTAS":
            continue

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions

        for celda in hoja[1]:
            celda.fill = fill_encabezado
            celda.font = Font(
                bold=True,
                color="FFFFFF",
            )
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for columna in hoja.columns:
            letra = columna[0].column_letter

            ancho = max(
                len(str(celda.value or ""))
                for celda in columna[:2000]
            )

            hoja.column_dimensions[letra].width = min(
                max(ancho + 2, 12),
                55,
            )

    # Resaltar columnas editables.
    columnas_editables = [
        "DECISION_MANUAL",
        "NIVEL_APROBADO",
        "PLAN_APROBADO",
        "CARRERA_APROBADA",
        "FUENTE_RESPALDO_DECISION",
        "FUNDAMENTO_DECISION",
        "RESPONSABLE_REVISION",
        "FECHA_REVISION",
        "ESTADO_CIERRE",
        "OBSERVACIONES_REVISOR",
    ]

    for nombre in columnas_editables:
        columna = encabezados[nombre]

        for fila in range(
            2,
            len(planilla) + 2,
        ):
            hoja_revision.cell(
                row=fila,
                column=columna,
            ).fill = fill_decision

    # Resaltar prioridad.
    col_prioridad = encabezados[
        "PRIORIDAD_REVISION"
    ]

    for fila in range(
        2,
        len(planilla) + 2,
    ):
        prioridad = hoja_revision.cell(
            row=fila,
            column=col_prioridad,
        ).value

        if prioridad == "P1_CRITICA":
            for celda in hoja_revision[fila]:
                celda.fill = fill_critico

            # Volver a marcar editables.
            for nombre in columnas_editables:
                columna = encabezados[nombre]
                hoja_revision.cell(
                    row=fila,
                    column=columna,
                ).fill = fill_decision

        elif prioridad == "P1_ALTA":
            hoja_revision.cell(
                row=fila,
                column=col_prioridad,
            ).fill = fill_alta


# ============================================================
# 12. MANIFIESTO
# ============================================================

manifiesto = {
    "fecha_ejecucion": datetime.now().isoformat(),
    "proceso": "Avance Curricular SIES 2026",
    "subproyecto": "5809 Matrícula",
    "anio_referencia_datos": 2025,
    "universo_revision_p1": 29,
    "distribucion": esperados,
    "decisiones_iniciales": 0,
    "aprobaciones_iniciales": 0,
    "fuentes_modificadas": False,
    "precargas_modificadas": False,
    "archivo_final_carga_generado": False,
    "excel_operativo": str(excel),
    "salida": str(SALIDA),
}

(SALIDA / "manifest_ejecucion.json").write_text(
    json.dumps(
        manifiesto,
        indent=2,
        ensure_ascii=False,
    ),
    encoding="utf-8",
)


# ============================================================
# 13. REPORTE
# ============================================================

lineas = [
    "PLANILLA OPERATIVA DE DECISIÓN — 29 CASOS P1",
    "=" * 96,
    "Proceso: Avance Curricular SIES 2026",
    "Subproyecto: 5809 Matrícula",
    "Año de referencia: 2025",
    "",
    "Composición:",
    "- Conflicto de carrera: 14",
    "- Múltiples trayectorias: 14",
    "- Nivel superior al máximo del plan: 1",
    "",
    "Las columnas amarillas son editables.",
    "No existen decisiones ni aprobaciones precargadas.",
    "",
    f"Excel: {excel}",
    f"Carpeta: {SALIDA}",
    "",
    "Fuentes modificadas: NO",
    "Precargas modificadas: NO",
    "Archivo final de carga generado: NO",
]

(REPORTES / "RESUMEN_EJECUCION.txt").write_text(
    "\n".join(lineas),
    encoding="utf-8",
)


# ============================================================
# 14. TERMINAL
# ============================================================

print()
print("=" * 100)
print("PLANILLA OPERATIVA DE DECISIÓN PARA LOS 29 CASOS P1")
print("=" * 100)
print(f"Casos incluidos: {len(planilla)}")
print()
print("COMPOSICIÓN")
print("-" * 100)

for _, fila in resumen_categorias.iterrows():
    print(
        f"{fila['CATEGORIA_REVISION']}: "
        f"{fila['CASOS']}"
    )

print()
print("Decisiones precargadas: 0")
print("Casos aprobados automáticamente: 0")
print("Estado inicial de todos los casos: PENDIENTE_REVISION")
print()
print(f"Excel operativo: {excel}")
print(f"Carpeta: {SALIDA}")
print("Fuentes modificadas: NO")
print("Precargas modificadas: NO")
print("Archivo final de carga generado: NO")
print("=" * 100)
