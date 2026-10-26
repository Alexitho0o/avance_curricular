"""Cierre agregado de auditoría de procedencia. No modifica bases ni asignaciones."""
from pathlib import Path
import csv, json, hashlib, subprocess, sqlite3, re
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

R=Path('/Users/alexi/Documents/GitHub/avance_curricular')
O=R/'outputs/origen_versiones_c3j4_e04_20261023_163300'
read=lambda p:json.loads(Path(p).read_text())
def sha(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()
def git(*args): return subprocess.check_output(['git','-C',str(R),*args],text=True).strip()
def norm(v): return str(int(v)) if isinstance(v,(int,float)) and v==int(v) else str(v or '').strip()
start=read(O/'INICIO.json'); c=read(O/'c3j4/PRIMARY.json')
e=read(O/'e04/regla_e04_preparacion.json'); em=read(O/'e04/regla_e04_manifest.json')
code=read(O/'BUSQUEDA_IMPLEMENTACIONES.json')
sources={f['archivo']:f['sha256'] for f in c['fuentes']}
sources.update(em['fuentes_sha256'])
sources[str(R/'scripts/construir_base_retencion_mu_2022_2026.py')]='edcade5f199443517ce6ec88a7c43a58cf52ee80a671ab0b9eddec0f1e6de19e'
sources[str(R/'codigo_gobernanza_v2.py')]=code['fallback_actual']['sha256']
frozen=dict(start['hashes_preservar'])
db=Path('/Users/alexi/.local/share/ipss/cierre_inf35_03_t3_20261023_154224/DICCIONARIO_SIES.sqlite')
frozen[str(db)]='625b9b9d449906b38be23b7c8ad6c7f4a762f43eab9f52f0bc71fe50a5cdb38e'
frozen[str(R/'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx')]='19cf33ae4925a7c13ec4745394141a297f4f64be4052a240cb0d2c42ffeb53bd'

# Relectura independiente de todos los localizadores de código, sin imprimir identidad.
checks=[]
for f in c['fuentes']:
    p=Path(f['archivo']); target=set(f['filas_foco']); codes=[]; formulas=0
    if p.suffix=='.csv':
        with p.open(encoding='utf-8-sig',newline='') as h:
            for rn,z in enumerate(csv.DictReader(h,delimiter=';'),2):
                if rn in target:
                    codes.append('I'+norm(z['COD_IES'])+'S'+norm(z['COD_SED'])+'C'+norm(z['COD_CAR'])+'J'+norm(z['JOR'])+'V'+norm(z['VERSION']))
    else:
        w=load_workbook(p,read_only=True,data_only=False);s=w[f['hoja']]
        hdr=list(next(s.iter_rows(values_only=True)))
        for rn,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
            if rn not in target:continue
            z=dict(zip(hdr,row));v=z.get('MC_CODIGO_UNICO')
            if v is None:v='I'+norm(z['COD_IES'])+'S'+norm(z['COD_SED'])+'C'+norm(z['COD_CAR'])+'J'+norm(z['JOR'])+'V'+norm(z['VERSION'])
            codes.append(v);formulas+=isinstance(z.get('MC_CODIGO_UNICO',z.get('VERSION')),str) and str(z.get('MC_CODIGO_UNICO',z.get('VERSION'))).startswith('=')
        w.close()
    checks.append({'archivo':str(p),'hoja':f['hoja'],'filas_verificadas':len(codes),'codigos':sorted(set(codes)),'ok':len(codes)==50 and set(codes)==set(f['foco_codigos']) and formulas==0})

ecells=[]
byfile={}
for rec in e['registros']:
    for col,val in rec['valores'].items():
        if col.endswith(':VERSION'):byfile.setdefault(rec['fuente'],[]).append((rec,col.split(':')[0]))
for file,rs in byfile.items():
    w=load_workbook(file,read_only=True,data_only=False)
    for rec,col in rs:
        cell=w[rec['hoja']][f"{col}{rec['fila']}"]
        ecells.append({'archivo':file,'hoja':rec['hoja'],'celda':cell.coordinate,'valor':cell.value,'tipo':cell.data_type,'ok':cell.value==1 and cell.data_type!='f'})
    w.close()
(O/'CONTRASTE_INDEPENDIENTE.json').write_text(json.dumps({'clasificacion':'B/C','c3j4':checks,'e04':ecells},ensure_ascii=False,indent=2)+'\n')

sourcechecks=[{'ruta':p,'sha256':sha(p),'esperado':h} for p,h in sources.items()]
frozenchecks=[{'ruta':p,'sha256':sha(p),'esperado':h} for p,h in frozen.items()]
state={'rama':git('branch','--show-current'),'HEAD':git('rev-parse','HEAD'),'remotos':git('remote','-v'),'status_short':git('status','--short'),'staged':git('diff','--cached','--name-only')}

report='''# Procedencia de versiones: C3/J4 2021–2022 y E04

**Resultado B/C:** los cambios de versión están reproducidos en los respaldos originales. **Resultado E:** no se localizó la consulta, fórmula o decisión que los generó. El punto de origen observable quedó acotado al lote de matrícula 2022 y a la precarga MU2025, respectivamente. No se adoptó ninguna equivalencia nueva.

## Alcance y separación

- **D:** solicitud de investigar 50 personas C3/J4 y E04, conservando bases, denominadores, diccionario y libro de retención. Sin commit, push ni consultas mutantes.
- **B:** proceso Matrícula Unificada; años de datos 2021–2022 y 2024–2026; ejecución 23-oct-2026. Repositorio principal `avance_curricular`, rama `feature/ire-2026`. No se está reconstruyendo una tasa ni el motor anual.
- **C:** identidad exacta normalizada; contexto de carrera/sede/jornada por año; fórmulas y valores literales inspeccionados separadamente. No se cruzó por nombre y no se atribuyó una asignación individual desde totales agregados.
- **E:** un código observado en MU acredita el valor allí registrado; no explica por sí solo por qué se seleccionó ni prueba la fecha efectiva de un cambio de plan. No se halló una nueva regla oficial A aplicable a la asignación investigada.

## 1. C3/J4: los 50 cambios

| Hallazgo | Conteo | Clase | Evidencia |
|---|---:|---|---|
| V1 en 2021 y V2 en 2022, mismo contexto atribuido IETMREOL20211/J4 en diccionario | 50 personas; 50 CODCLI | B/C | c3j4/PRIMARY.json: seis fuentes, hojas, filas y hashes |
| VERSION/código único escritos como valores; ninguna fórmula en las celdas revisadas | 50 por cada fuente | B | PRIMARY.json; CONTRASTE_INDEPENDIENTE.json |
| FECH_CAR de los 50 cambios | 06/05/2022 16:14 | B | RAW 2022.csv; PRIMARY.json |
| Cambia también el año de origen desde un año anterior a 2021 hacia 2021 | 11 de los 50 | B | c3j4/PATRONES.json; no se atribuye causalidad |

**B/C — Control de contexto, no ampliación del denominador:** entre las 75 personas presentes en la familia C3/J4 en ambos años hay 58 V1→V2 (incluyen las 50), 15 V2→V2 y 2 V1→V1. Las 58 cambiadas y las 15 que ya tenían V2 pertenecen al lote 16:14; las dos que conservan V1, al lote 00:11 del mismo día. En la familia completa RAW2022 hay tres filas V1 del lote 00:11 y 155 V2: 153 del lote 16:14 y dos del lote 16:38.

**B — Las variables académicas contrastadas no explican por sí solas la separación:** las dos personas que conservan V1 comparten con las 50 foco ingreso actual 2021-1, modalidad 3, FOR_ING_ACT=1, VIG=1, perfil regular de ocho semestres y contexto IETMREOL20211 en el diccionario. Los informes primarios no conservan CODCLI/PLAN_DE_ESTUDIO; este contexto del diccionario no demuestra automáticamente el plan efectivo en la fecha histórica.

**C — La consolidación no asignó estas versiones:** `scripts/construir_base_retencion_mu_2022_2026.py`, líneas 227–265, lee la fuente, mapea la columna y copia el valor mediante `typed_value`. VERSION ya era V2 en RAW2022. Se verificaron 50/50 en las seis fuentes sin colapso de múltiples códigos: 75 personas comunes tienen una fila/código por familia en cada año. RAW2022 declara TIPO_DOC=R; MAT2021 no contiene ese campo y se usó MM_N_DOCUMENTO+MM_DV, sin inventarlo.

**E — Pendiente exacto:** falta el extracto de preparación anterior al lote 06/05/2022 16:14, su consulta/script o un registro de decisión que explique el cambio a V2 y la conservación de V1 en el lote 00:11. FECH_CAR separa los lotes observados, pero no demuestra migración, corrección ni error. No corresponde convertir esa fecha en una regla de asignación.

## 2. E04: V2→V1→V2

**B — La secuencia documentada permanece:** I162S2C1J1V2 en 2024, I162S2C1J1V1 en 2025 e I162S2C1J1V2 en 2026. E04 se identificó de forma independiente y no es INF35-03.

| Punto observado en respaldo MU2025 | Valor | Clase |
|---|---:|---|
| IPSS_Procesamiento.xlsx · Precarga!Q1025 | V1 literal | B |
| Mismo libro · Carga_Completa!Q1027; Carga_Completa (2)/(3)!Q1032; Final a carga!Q1004 | V1 literal | B |
| IPSS2025.xlsx · O1004; descarga de lo cargado · Q1020; RAW2025 · AA867 | V1 literal | B |
| FECH_CAR de E04 en la precarga | 16-abr-2025 11:50:07 | B |

**B/C:** 17 libros de preparación, 32 filas coincidentes y 10 celdas VERSION=1, todas literales. Incluyen copias: no son 32 evidencias independientes ni demuestran el orden de edición entre hojas. No hay fórmula de VERSION ni conexiones/queryTables/VBA en el libro de preparación que expliquen esa asignación. El incremental denominado 02-01-2025 no contiene a E04; el RAW2024 de esa carpeta sí contiene V2.

**B — Control de contexto:** Precarga C1/J1/S2 contiene 22 filas: 14 V2 y ocho V1. Las ocho V1 tienen FECH_CAR del 16 de abril de 2025; también hay una V2 de ese día. De las ocho personas con código de la misma familia en RAW2024, siete mantienen V2 y E04 pasa a V1. Ni el día de carga ni la sola familia bastan para justificar V1.

**E — Pendiente exacto:** la asignación V1 ya venía incorporada en `Precarga!Q1025`. Falta la consulta, transformación o decisión previa a esa precarga/lote. FECH_CAR no acredita cuándo cambió el plan ni quién decidió la versión. Se conserva la secuencia anual observada; no se fuerza una V2 retrospectiva.

## 3. Scripts e historial local

**C:** se inventariaron las referencias Git locales de avance_curricular, CARACTERIZACIÓN e IPSS_Procesos_SIES, sin fetch. Se inspeccionaron estáticamente 756, 2.590 y 143 blobs de extensiones de código, respectivamente (incluyen dependencias); 67 contienen expresiones candidatas. Las coincidencias de texto son una búsqueda, no 67 reglas validadas. `BUSQUEDA_IMPLEMENTACIONES.json` conserva objetos Git, rutas, hashes y líneas. De 150 rutas de código del inventario OneDrive previo, se releyeron seis ajenas a FCU2026; 144 respaldos FCU2026 se conservaron como inventariados fuera de alcance, no como revisados para MU2022/2025.

**C — Implementación encontrada, alcance distinto:** `codigo_gobernanza_v2.py`, líneas 3763–3781, primero extrae VERSION del código final; si falta y hay COD_CAR+JOR, usa la versión máxima del catálogo por esos campos. El propio generador referencia MU/oferta 2026. Su aparición en Git figura en commits d0dc2ebd y 5ddb6868. Esto no acredita ejecución sobre los archivos de 2022 o 2025 ni explica esos lotes. No se ejecutó ni alteró ese motor.

**C:** los blobs históricos `scripts/auditar_logica_versiones_sies_mu.py` (7e6f54e2), `scripts/reconstruir_rectificacion_mu2026_desde_cero.py` (053d4935) y `scripts/sies_resolver.py` (2ad3ccf1) son auditoría, rectificación MU2026 y validación/parseo de códigos. No se localizó en ellos una instrucción que produjera el lote2022 o la precarga2025 del foco. Las copias y las fechas de los archivos tampoco prueban autoría o ejecución.

**E:** no localizar el generador en esta cobertura no equivale a afirmar que nunca existió. No se ha demostrado una regla universal basada en ingreso, plan, lote o versión máxima.

## 4. Resultado operativo

**D:** no aplicar cierres ni cambios de código con esta investigación. Los códigos por año ya estaban documentados; ahora se conoce mejor el tramo donde se pierde la explicación de su asignación. Las dos preguntas causales siguen E.

**B/C:** QA de cardinalidad, relectura de celdas, hashes, privacidad y Git en `QA.json`; evidencia original con rutas/hojas/filas/hashes en los manifiestos de cada subcarpeta y `MANIFEST.json`. Se mantienen Excel, diccionario, bases y denominadores. Ningún dato individual se publica ni incorpora a estos resultados agregados. No se hizo commit ni push.

**E — Reanudación precisa:** buscar el insumo/consulta anterior al CSV2022 del lote16:14 y el insumo/consulta anterior a Precarga2025 con FECH_CAR16-abr. No volver a probar reglas por mayoría o por año sobre estos casos, ni presentar coincidencia de lotes como explicación causal.
'''
(O/'INFORME.md').write_text(report)
checkpoint={'fase':'3/3','hito':'2/2','estado':'Investigación de procedencia completada; dos reglas generadoras no localizadas (E)',
 'alcance':'50 C3/J4 y E04','cambios_adoptados':0,'modifica_denominadores':False,
 'siguiente_evidencia':['Consulta/insumo anterior a RAW2022 FECH_CAR06/05/2022 16:14','Consulta/insumo anterior a Precarga MU2025!Q1025 FECH_CAR16/04/2025 11:50:07'],
 'no_repetir':'Códigos históricos y lotes ya contrastados; no inferir asignación desde totales, mayoría o año.',
 'git_antes':{'rama':start['rama'],'HEAD':start['HEAD']},'git_despues':state}
(O/'CHECKPOINT.json').write_text(json.dumps(checkpoint,ensure_ascii=False,indent=2)+'\n')

# Privacidad: comprobar ausencia de identificadores conocidos en las salidas nuevas.
conn=sqlite3.connect('file:'+str(db)+'?mode=ro',uri=True)
ids=set()
for (payload,) in conn.execute('select payload from observaciones'):
    row=json.loads(payload)
    ids.update(str(row[k]) for k in ['persona','CODCLI'] if row.get(k))
conn.close()
pat=re.compile(r'(?<![A-Za-z0-9])(?:'+'|'.join(re.escape(v) for v in sorted(ids,key=len,reverse=True))+r')(?![A-Za-z0-9])')
files=list(O.rglob('*'))+list((R/'scripts/origen_versiones_c3j4_e04_20261023').rglob('*.py'))+[Path(__file__),R/'scripts/buscar_asignacion_version_mu_20261023.py',R/'scripts/inventariar_reglas_mu_c3j4_e04_20261023.py',R/'docs/ESPECIFICACION_ORIGEN_VERSIONES_C3J4_E04_20261023.md']
leaks=[str(p) for p in files if p.is_file() and p.suffix in ['.md','.json','.py'] and pat.search(p.read_text())]
qa={'clasificacion':'B/C','controles':{
 '50_personas_fuentes_reproducidas':len(checks)==6 and all(q['ok'] for q in checks),
 'E04_10_celdas_literal_V1':len(ecells)==10 and all(q['ok'] for q in ecells),
 'QA_C3J4':read(O/'c3j4/QA.json')['estado']=='PASS',
 'QA_E04':read(O/'e04/regla_e04_QA.json')['fuentes_intactas'],
 'fuentes_hash_intactos':all(x['sha256']==x['esperado'] for x in sourcechecks),
 'base_excel_diccionario_intactos':all(x['sha256']==x['esperado'] for x in frozenchecks),
 'rama_HEAD_sin_cambio':state['rama']==start['rama'] and state['HEAD']==start['HEAD'],
 'sin_archivos_stage':not state['staged'],
 'sin_identificadores_conocidos':not leaks},
 'fuentes':sourcechecks,'preservados':frozenchecks,'hallazgos_privacidad_rutas':leaks,
 'limite':'Valida el rastreo y conservación; no valida una regla causal no localizada ni acredita fechas de cambios de plan.'}
qa['estado']='PASS' if all(qa['controles'].values()) else 'FAIL'
(O/'QA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
assert qa['estado']=='PASS',qa['controles']
manifest={'fecha':datetime.now(ZoneInfo('America/Santiago')).isoformat(),'proceso':'Matrícula Unificada / procedencia de VERSION SIES','clasificaciones':{'B':'valores y lotes observados','C':'rastreo/inspección estática/relectura','D':'preservación sin adopciones','E':'causa de dos asignaciones pendiente'},'fuentes':sources,'preservados':frozen,
 'scripts':{str(p):sha(p) for p in files if p.is_file() and p.suffix=='.py'},
 'salidas':{str(p.relative_to(R)):sha(p) for p in O.rglob('*') if p.is_file() and p.name!='MANIFEST.json'},
 'git_antes':{'rama':start['rama'],'HEAD':start['HEAD']},'git_despues':state,'QA':qa['estado'],'commit':False,'push':False,'adjudicaciones':0}
(O/'MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'QA':qa['estado'],'controles':qa['controles'],'fuentes':len(sourcechecks),'informe':str(O/'INFORME.md'),'hash_informe':sha(O/'INFORME.md'),'hash_manifest':sha(O/'MANIFEST.json'),'HEAD':state['HEAD']},ensure_ascii=False))
