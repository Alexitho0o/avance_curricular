#!/usr/bin/env python3

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular").resolve()
EXEC_DEFAULT = (
    REPO
    / "avance_curricular_2026"
    / "05_cierre_integral"
    / "CIERRE_AVANCE_CURRICULAR_20260703_125157"
)
PREP_DEFAULT = (
    EXEC_DEFAULT
    / "04_CONCILIACION_MATRICULA"
    / "PREPARACION_MATRICULA_20260703_141254"
)

SUBDIRS = [
    "00_CONTROL",
    "01_INVENTARIO",
    "02_PERFIL_FUENTES",
    "03_MAPEO_CAMPOS",
    "04_COBERTURA",
    "05_IDENTIDAD",
    "06_CONCILIACION",
    "07_BRECHAS",
    "08_SOLICITUDES",
    "09_AUDITORIA",
    "10_REPORTES",
]

ALLOWED_STATUS = {
    "NO_INICIADA",
    "EN_EJECUCION",
    "COMPLETADA",
    "COMPLETADA_CON_BRECHAS",
    "BLOQUEADA",
    "FALLIDA",
}

SCAN_EXTENSIONS = {
    ".xlsx",
    ".xlsm",
    ".csv",
    ".tsv",
    ".txt",
    ".json",
    ".py",
    ".md",
}

SKIP_DIRS = {
    ".git",
    ".venv",
    "__pycache__",
    ".mypy_cache",
    ".pytest_cache",
}

CATEGORIES = {
    "RUT": ["RUT", "RUN", "DOCUMENTO", "NUM_DOCUMENTO", "N_DOC"],
    "CODCLI": ["CODCLI"],
    "CARRERA": ["CODCARR", "CODCARPR", "CODIGO_UNICO", "CARRERA"],
    "PLAN": ["CODPESTUD", "PLAN_DE_ESTUDIO", "PLAN_ESTUDIOS", "NOMPESTUD"],
    "PERIODO": ["ANO", "AÑO", "PERIODO", "SEMESTRE", "ANOMATRICULA", "PERIODOMATRICULA"],
    "NIVEL": ["NIVEL"],
    "ASIGNATURA": ["CODRAMO", "RAMOEQUIV", "ASIGNATURA", "RAMO", "UNIDAD"],
    "NOTA": ["NOTA", "NOTA_FINAL", "EXAMEN"],
    "ESTADO": ["ESTADO", "DESCRIPCION_ESTADO", "SITUACION"],
    "APROBACION": ["APROBADO", "REPROBADO", "APROBADA", "APROBACION"],
    "CREDITOS": ["CREDITO", "CREDITOS", "SCT", "HORAS", "HOR_TEO", "HOR_PRAC", "MODULOS"],
    "CONVALIDACION": ["CONVALIDADO", "CONVALIDACION", "CONVALIDADA", "HOMOLOG", "RECONOC", "VALIDACION"],
    "INSCRIPCION": ["INSCRIPCION", "INSCRITA", "INSCRITO", "CARGA", "MATRICULA"],
}

SIES_FIELDS = [
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
    "PLAN_ESTUDIOS",
    "VIGENCIA",
]


def now_iso() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm(value: Any) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip().upper()
    if re.fullmatch(r"-?\d+\.0", text):
        text = text[:-2]
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", "_", text).strip("_")


def compact(value: Any) -> str:
    return re.sub(r"[^0-9K]", "", norm(value))


def split_rut(value: Any) -> tuple[str, str]:
    text = str(value or "").strip()
    if "-" in text:
        body, dv = text.rsplit("-", 1)
        return compact(body), compact(dv)
    clean = compact(text)
    if len(clean) > 1 and clean[-1] in "0123456789K":
        return clean[:-1], clean[-1]
    return clean, ""


def join_unique(values: Any, limit: int | None = None) -> str:
    seen: set[str] = set()
    out: list[str] = []
    for value in list(values):
        text = "" if pd.isna(value) else str(value).strip()
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
    out = sorted(out)
    if limit is not None and len(out) > limit:
        return " | ".join(out[:limit]) + f" | ...(+{len(out)-limit})"
    return " | ".join(out)


def detect_encoding(path: Path) -> str:
    data = path.read_bytes()[:200000]
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            data.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def detect_delimiter(text: str) -> str:
    first = next((line for line in text.splitlines() if line.strip()), "")
    counts = {sep: first.count(sep) for sep in ("\t", ";", ",", "|")}
    sep = max(counts, key=counts.get)
    return sep if counts[sep] > 0 else ""


def classify_path(path: Path) -> str:
    s = str(path).lower()
    name = path.name.lower()
    if "precarga" in name:
        return "PRECARGA"
    if "fuentes_institucionales" in s or "/input/" in s:
        return "FUENTE_INSTITUCIONAL"
    if "avance" in name and ("malla" in name or "individual" in name):
        return "REPORTE_INDIVIDUAL"
    if "resultado" in s or "resultados" in s:
        return "RESULTADO_PREVIO"
    if "auditoria" in s or "auditorias" in s:
        return "AUDITORIA"
    if path.suffix.lower() == ".py":
        return "SCRIPT"
    if "evidencia" in s:
        return "EVIDENCIA"
    if "trabajo" in s:
        return "ARCHIVO_DE_TRABAJO"
    return "FUENTE_NO_CLASIFICADA"


def header_categories(columns: list[str]) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    normalized = {col: norm(col) for col in columns}
    for category, patterns in CATEGORIES.items():
        hits = []
        for col, ncol in normalized.items():
            if any(pat in ncol for pat in patterns):
                hits.append(col)
        if hits:
            result[category] = hits
    return result


def infer_granularity(cats: dict[str, list[str]]) -> str:
    has_student = bool(cats.get("RUT") or cats.get("CODCLI"))
    if has_student and cats.get("ASIGNATURA"):
        return "UNA_FILA_POR_ESTUDIANTE_ASIGNATURA"
    if has_student and cats.get("PERIODO"):
        return "UNA_FILA_POR_ESTUDIANTE_PERIODO"
    if has_student:
        return "UNA_FILA_POR_ESTUDIANTE_O_TRAYECTORIA"
    if cats.get("PLAN") and cats.get("ASIGNATURA"):
        return "UNA_FILA_POR_PLAN_ASIGNATURA"
    if cats.get("PLAN"):
        return "UNA_FILA_POR_PLAN"
    if cats.get("ASIGNATURA"):
        return "UNA_FILA_POR_CURSO"
    return "NO_DETERMINADA"


def possible_use(cats: dict[str, list[str]]) -> str:
    uses = []
    if cats.get("RUT") or cats.get("CODCLI"):
        uses.append("IDENTIDAD")
    if cats.get("PERIODO") and cats.get("ASIGNATURA"):
        uses.append("AVANCE_ANUAL")
    if cats.get("ESTADO") or cats.get("APROBACION") or cats.get("NOTA"):
        uses.append("APROBACION")
    if cats.get("CREDITOS"):
        uses.append("UNIDADES")
    if cats.get("CONVALIDACION"):
        uses.append("CONVALIDACIONES_RECONOCIMIENTOS")
    if cats.get("PLAN"):
        uses.append("PLAN")
    if cats.get("CARRERA"):
        uses.append("CARRERA_TRAYECTORIA")
    return " | ".join(uses) if uses else "NO_DETERMINADO"


def is_useful(cats: dict[str, list[str]]) -> bool:
    has_identity = bool(cats.get("RUT") or cats.get("CODCLI"))
    has_academic = bool(
        cats.get("PERIODO")
        or cats.get("ASIGNATURA")
        or cats.get("ESTADO")
        or cats.get("NOTA")
        or cats.get("CREDITOS")
        or cats.get("CONVALIDACION")
    )
    return has_identity and has_academic


def sample_csv(path: Path) -> tuple[list[str], pd.DataFrame, str, str]:
    enc = detect_encoding(path)
    raw = path.read_bytes()[:300000]
    text = raw.decode(enc, errors="replace")
    sep = detect_delimiter(text)
    if not sep:
        return [], pd.DataFrame(), enc, ""
    df = pd.read_csv(path, sep=sep, encoding=enc, dtype=str, keep_default_na=False, nrows=500)
    return list(df.columns), df, enc, sep


def workbook_sheets(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        headers = []
        sample_rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            if i == 1:
                headers = values
            elif i <= 25:
                sample_rows.append(values)
            else:
                break
        sheets.append(
            {
                "sheet": ws.title,
                "rows": ws.max_row,
                "columns": ws.max_column,
                "headers": headers,
                "sample_rows": sample_rows,
            }
        )
    return sheets


def profile_text_like(path: Path) -> tuple[list[str], str]:
    enc = detect_encoding(path)
    text = path.read_bytes()[:300000].decode(enc, errors="replace")
    lines = text.splitlines()
    headers = []
    if lines:
        sep = detect_delimiter(lines[0])
        if sep:
            headers = lines[0].split(sep)
        else:
            headers = re.findall(r"[A-Za-z_ÁÉÍÓÚáéíóúÑñ0-9]{3,}", "\n".join(lines[:80]))
    return headers[:200], enc


def write_tsv(path: Path, df: pd.DataFrame) -> None:
    df.to_csv(path, sep="\t", index=False)


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def write_xlsx(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            safe = re.sub(r"[\[\]\*\?/\\:]", "_", name)[:31]
            df.to_excel(writer, sheet_name=safe, index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                letter = col[0].column_letter
                width = max(len(str(cell.value or "")) for cell in col[:500])
                ws.column_dimensions[letter].width = min(max(width + 2, 12), 60)


def append_tsv(path: Path, row: dict[str, Any], columns: list[str]) -> None:
    exists = path.exists() and path.stat().st_size > 0
    with path.open("a", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        if not exists:
            writer.writeheader()
        writer.writerow(row)


def init_run(exec_dir: Path, run_dir: Path | None) -> Path:
    if run_dir is None:
        run_dir = exec_dir / "05_RESOLUCION_PENDIENTES" / f"DIAGNOSTICO_FUENTES_AVANCE_{ts()}"
    for sub in SUBDIRS:
        (run_dir / sub).mkdir(parents=True, exist_ok=True)
    return run_dir.resolve()


def ctl(run_dir: Path) -> dict[str, Path]:
    c = run_dir / "00_CONTROL"
    return {
        "estado": c / "ESTADO_DIAGNOSTICO_FUENTES.json",
        "checkpoint": c / "CHECKPOINT.json",
        "manifest_fuentes": c / "MANIFEST_FUENTES.json",
        "manifest_salidas": c / "MANIFEST_SALIDAS.json",
        "bloqueos": c / "BLOQUEOS.tsv",
        "decisiones": c / "DECISIONES.tsv",
        "contradicciones": c / "CONTRADICCIONES.tsv",
        "validaciones": c / "VALIDACIONES.tsv",
        "reanudar": c / "REANUDAR_DESDE_AQUI.md",
    }


def read_5809_from_prep(prep_dir: Path) -> pd.DataFrame:
    df = pd.read_excel(prep_dir / "01_UNIVERSO" / "UNIVERSO_MATRICULA.xlsx", sheet_name="UNIVERSO_5809", dtype=str, keep_default_na=False)
    if len(df) != 2371:
        raise RuntimeError(f"Universo 5809 cambió: {len(df)}")
    return df


def read_identity(prep_dir: Path) -> pd.DataFrame:
    return pd.read_csv(prep_dir / "02_IDENTIFICADORES" / "MAPEO_FILA_5809_IDENTIDAD.tsv", sep="\t", dtype=str, keep_default_na=False)


def load_source_dataframe(source_id: str, ruta: str, hoja: str = "") -> pd.DataFrame:
    path = Path(ruta)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return pd.read_excel(path, sheet_name=hoja, dtype=str, keep_default_na=False)
    sep = "\t" if path.suffix.lower() == ".tsv" else None
    if sep is None:
        _, _, _, sep = sample_csv(path)
    return pd.read_csv(path, sep=sep, dtype=str, keep_default_na=False)


class Runner:
    def __init__(self, exec_dir: Path, prep_dir: Path, run_dir: Path):
        self.exec_dir = exec_dir
        self.prep_dir = prep_dir
        self.run_dir = run_dir
        self.paths = ctl(run_dir)
        self.state = {
            "proceso": "Avance Curricular SIES 2026",
            "subproceso": "Diagnostico fuentes de avance",
            "fecha_inicio": now_iso(),
            "fecha_ultima_actualizacion": now_iso(),
            "estado": "EN_EJECUCION",
            "fase_actual": "A0",
            "fases": {f"A{i}": "NO_INICIADA" for i in range(11)},
            "csv_generado": False,
            "sies_ready_generado": False,
            "carga_pes_realizada": False,
        }
        self.inventory = pd.DataFrame()
        self.profile = pd.DataFrame()
        self.useful = pd.DataFrame()
        self.df5809 = pd.DataFrame()
        self.identity = pd.DataFrame()

    def mark(self, phase: str, status: str) -> None:
        if status not in ALLOWED_STATUS and not status.startswith("DIAGNOSTICO_FUENTES"):
            raise RuntimeError(status)
        self.state["fase_actual"] = phase
        self.state["fases"][phase] = status
        self.state["fecha_ultima_actualizacion"] = now_iso()
        if status.startswith("DIAGNOSTICO_FUENTES"):
            self.state["estado"] = status
        elif status == "BLOQUEADA":
            self.state["estado"] = "BLOQUEADO_POR_AUSENCIA_DE_FUENTE_AVANCE"
        else:
            self.state["estado"] = "EN_EJECUCION"
        write_json(self.paths["estado"], self.state)

    def validation(self, phase: str, validation: str, status: str, detail: str) -> None:
        append_tsv(
            self.paths["validaciones"],
            {"FASE": phase, "VALIDACION": validation, "ESTADO": status, "DETALLE": detail},
            ["FASE", "VALIDACION", "ESTADO", "DETALLE"],
        )

    def decision(self, phase: str, decision: str, evidence: str, impact: str) -> None:
        append_tsv(
            self.paths["decisiones"],
            {"FECHA": now_iso(), "FASE": phase, "DECISION": decision, "EVIDENCIA": evidence, "IMPACTO": impact},
            ["FECHA", "FASE", "DECISION", "EVIDENCIA", "IMPACTO"],
        )

    def block(self, phase: str, kind: str, detail: str, action: str) -> None:
        append_tsv(
            self.paths["bloqueos"],
            {"FECHA": now_iso(), "FASE": phase, "TIPO": kind, "DETALLE": detail, "ACCION_REQUERIDA": action},
            ["FECHA", "FASE", "TIPO", "DETALLE", "ACCION_REQUERIDA"],
        )

    def checkpoint(self, phase: str, status: str, outputs: list[Path], next_phase: str) -> None:
        write_json(
            self.paths["checkpoint"],
            {
                "fase": phase,
                "estado": status,
                "fecha": now_iso(),
                "salidas": [str(path) for path in outputs],
                "siguiente_fase": next_phase,
            },
        )
        command = (
            f'python3 scripts/diagnosticar_fuentes_avance_sies2026.py '
            f'--execution-dir "{self.exec_dir}" --prep-dir "{self.prep_dir}" --run-dir "{self.run_dir}"'
        )
        self.paths["reanudar"].write_text(
            "\n".join(
                [
                    "# Reanudar diagnostico de fuentes de avance",
                    "",
                    f"Estado: {status}",
                    f"Carpeta: {self.run_dir}",
                    "",
                    "```bash",
                    command,
                    "```",
                    "",
                ]
            ),
            encoding="utf-8",
        )

    def a0(self) -> None:
        self.mark("A0", "EN_EJECUCION")
        shutil.copy2(Path(__file__), self.run_dir / "09_AUDITORIA" / Path(__file__).name)
        required = [
            self.exec_dir / "00_CONTROL" / "ESTADO_EJECUCION.json",
            self.prep_dir / "00_CONTROL" / "ESTADO_PREPARACION_MATRICULA.json",
            self.prep_dir / "01_UNIVERSO" / "UNIVERSO_MATRICULA.xlsx",
            self.prep_dir / "02_IDENTIFICADORES" / "MAPEO_FILA_5809_IDENTIDAD.tsv",
        ]
        missing = [str(path) for path in required if not path.exists()]
        if missing:
            self.block("A0", "FALTA_FUENTE_OBLIGATORIA", " | ".join(missing), "Restituir artefactos de cierre/preparacion.")
            self.mark("A0", "BLOQUEADA")
            raise RuntimeError("Faltan artefactos obligatorios.")
        self.df5809 = read_5809_from_prep(self.prep_dir)
        self.identity = read_identity(self.prep_dir)
        self.validation("A0", "CARPETA_DIAGNOSTICO_CREADA", "OK", str(self.run_dir))
        self.validation("A0", "UNIVERSO_5809_VERIFICADO", "OK", "2371")
        self.mark("A0", "COMPLETADA")
        self.checkpoint("A0", "COMPLETADA", [self.paths["estado"]], "A1")

    def scan_files(self) -> list[Path]:
        files: list[Path] = []
        for root, dirs, names in os.walk(REPO):
            root_path = Path(root)
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
            # Evita re-inventariar la ejecucion que se esta construyendo.
            if self.run_dir in [root_path, *root_path.parents]:
                continue
            for name in names:
                if name.startswith("~$") or name == ".DS_Store":
                    continue
                path = root_path / name
                if path.suffix.lower() in SCAN_EXTENSIONS:
                    files.append(path)
        return sorted(files, key=lambda p: str(p))

    def profile_path(self, source_id: str, path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        stat = path.stat()
        base = {
            "ID_FUENTE": source_id,
            "NOMBRE": path.name,
            "RUTA": str(path),
            "CLASIFICACION": classify_path(path),
            "ORIGEN": "REPOSITORIO",
            "TAMANO_BYTES": stat.st_size,
            "FECHA_MODIFICACION": datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds"),
            "SHA256": sha256(path),
            "FORMATO": path.suffix.lower().lstrip("."),
        }
        rows: list[dict[str, Any]] = []
        aggregate = {"content_score": 0, "name_score": 0, "estado": "DESCARTADO_POR_CONTENIDO"}
        name_norm = norm(path.name)
        aggregate["name_score"] = sum(1 for words in CATEGORIES.values() for word in words if word in name_norm)
        try:
            if path.suffix.lower() in {".xlsx", ".xlsm"}:
                sheets = workbook_sheets(path)
                for sheet in sheets:
                    cats = header_categories(sheet["headers"])
                    content_score = sum(len(v) for v in cats.values())
                    aggregate["content_score"] += content_score
                    row = {
                        **base,
                        "HOJAS": len(sheets),
                        "HOJA": sheet["sheet"],
                        "FILAS": sheet["rows"],
                        "COLUMNAS": sheet["columns"],
                        "ENCABEZADOS": " | ".join(sheet["headers"][:80]),
                        "CODIFICACION": "",
                        "DELIMITADOR": "",
                        "CATEGORIAS_DETECTADAS": " | ".join(sorted(cats.keys())),
                        "GRANULARIDAD": infer_granularity(cats),
                        "POSIBLE_USO": possible_use(cats),
                        "ESTADO_VALIDACION": "CANDIDATA_POR_CONTENIDO" if is_useful(cats) else "DESCARTADA_O_SECUNDARIA",
                    }
                    rows.append(row)
            elif path.suffix.lower() in {".csv", ".tsv"}:
                headers, sample, enc, sep = sample_csv(path)
                cats = header_categories(headers)
                aggregate["content_score"] += sum(len(v) for v in cats.values())
                rows.append(
                    {
                        **base,
                        "HOJAS": "",
                        "HOJA": "",
                        "FILAS": "",
                        "COLUMNAS": len(headers),
                        "ENCABEZADOS": " | ".join(headers[:80]),
                        "CODIFICACION": enc,
                        "DELIMITADOR": sep,
                        "CATEGORIAS_DETECTADAS": " | ".join(sorted(cats.keys())),
                        "GRANULARIDAD": infer_granularity(cats),
                        "POSIBLE_USO": possible_use(cats),
                        "ESTADO_VALIDACION": "CANDIDATA_POR_CONTENIDO" if is_useful(cats) else "DESCARTADA_O_SECUNDARIA",
                    }
                )
            else:
                headers, enc = profile_text_like(path)
                cats = header_categories(headers)
                aggregate["content_score"] += sum(len(v) for v in cats.values())
                rows.append(
                    {
                        **base,
                        "HOJAS": "",
                        "HOJA": "",
                        "FILAS": "",
                        "COLUMNAS": len(headers),
                        "ENCABEZADOS": " | ".join(headers[:80]),
                        "CODIFICACION": enc,
                        "DELIMITADOR": "",
                        "CATEGORIAS_DETECTADAS": " | ".join(sorted(cats.keys())),
                        "GRANULARIDAD": "SCRIPT_O_TEXTO" if path.suffix.lower() in {".py", ".md", ".txt"} else "NO_DETERMINADA",
                        "POSIBLE_USO": possible_use(cats),
                        "ESTADO_VALIDACION": "CANDIDATA_POR_CONTENIDO" if is_useful(cats) else "DESCARTADA_O_SECUNDARIA",
                    }
                )
        except Exception as exc:
            rows.append(
                {
                    **base,
                    "HOJAS": "",
                    "HOJA": "",
                    "FILAS": "",
                    "COLUMNAS": "",
                    "ENCABEZADOS": "",
                    "CODIFICACION": "",
                    "DELIMITADOR": "",
                    "CATEGORIAS_DETECTADAS": "",
                    "GRANULARIDAD": "",
                    "POSIBLE_USO": "",
                    "ESTADO_VALIDACION": f"PENDIENTE_CLASIFICACION_ERROR: {type(exc).__name__}: {exc}",
                }
            )
        return rows, aggregate

    def a1(self) -> None:
        self.mark("A1", "EN_EJECUCION")
        rows = []
        for idx, path in enumerate(self.scan_files(), start=1):
            prof_rows, _ = self.profile_path(f"FAV{idx:05d}", path)
            rows.extend(prof_rows)
        inv = pd.DataFrame(rows)
        if inv.empty:
            self.block("A1", "SIN_FUENTES_ESCANEADAS", "No se encontraron archivos con extensiones objetivo.", "Revisar repositorio.")
            self.mark("A1", "BLOQUEADA")
            raise RuntimeError("Sin fuentes escaneadas.")
        useful = inv[inv["ESTADO_VALIDACION"].eq("CANDIDATA_POR_CONTENIDO")].copy()
        discarded = inv[~inv["ESTADO_VALIDACION"].eq("CANDIDATA_POR_CONTENIDO")].copy()
        pending = inv[inv["ESTADO_VALIDACION"].astype(str).str.startswith("PENDIENTE")].copy()
        hashes = inv[["ID_FUENTE", "RUTA", "SHA256"]].drop_duplicates()
        self.inventory = inv
        self.useful = useful
        write_tsv(self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.tsv", inv)
        write_xlsx(
            self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.xlsx",
            {
                "INVENTARIO": inv,
                "CANDIDATAS": useful,
                "DESCARTADAS": discarded.head(20000),
                "PENDIENTES": pending,
            },
        )
        write_tsv(self.run_dir / "01_INVENTARIO" / "HASHES_FUENTES.tsv", hashes)
        write_tsv(self.run_dir / "01_INVENTARIO" / "ARCHIVOS_DESCARTADOS.tsv", discarded)
        write_tsv(self.run_dir / "01_INVENTARIO" / "ARCHIVOS_PENDIENTES_CLASIFICACION.tsv", pending)
        write_json(
            self.paths["manifest_fuentes"],
            {
                "fecha": now_iso(),
                "fuentes_inventariadas_tablas_hojas": len(inv),
                "fuentes_candidatas_tablas_hojas": len(useful),
                "criterio": "Candidata solo si la estructura inspeccionada contiene identificadores y variables academicas.",
            },
        )
        self.validation("A1", "INVENTARIO_DIRIGIDO_COMPLETO", "OK", f"filas inventario={len(inv)} candidatas={len(useful)}")
        if useful.empty:
            self.block("A1", "SIN_FUENTE_AVANCE_ADICIONAL", "No se encontraron fuentes candidatas por contenido.", "Preparar expediente formal de fuente faltante.")
        self.mark("A1", "COMPLETADA" if not useful.empty else "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A1", "COMPLETADA", [self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.xlsx"], "A2")

    def a2(self) -> None:
        self.mark("A2", "EN_EJECUCION")
        inv = self.inventory if not self.inventory.empty else pd.read_csv(self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.tsv", sep="\t", dtype=str, keep_default_na=False)
        candidates = inv[inv["ESTADO_VALIDACION"].eq("CANDIDATA_POR_CONTENIDO")].copy()
        dict_rows = []
        catalog_rows = []
        temporal_rows = []
        granular_rows = []
        for _, row in candidates.iterrows():
            headers = [h for h in str(row["ENCABEZADOS"]).split(" | ") if h]
            cats = header_categories(headers)
            for col in headers:
                col_cats = [cat for cat, values in cats.items() if col in values]
                if col_cats:
                    dict_rows.append(
                        {
                            "ID_FUENTE": row["ID_FUENTE"],
                            "RUTA": row["RUTA"],
                            "HOJA": row["HOJA"],
                            "COLUMNA": col,
                            "CATEGORIAS": " | ".join(col_cats),
                        }
                    )
            temporal_rows.append(
                {
                    "ID_FUENTE": row["ID_FUENTE"],
                    "RUTA": row["RUTA"],
                    "HOJA": row["HOJA"],
                    "RANGO_TEMPORAL_OBSERVADO": "CONTIENE_COLUMNA_ANO_PERIODO" if "PERIODO" in str(row["CATEGORIAS_DETECTADAS"]) else "NO_DETERMINADO_EN_PERFIL",
                    "CONTIENE_2025": "PENDIENTE_CONFIRMACION_MUESTRA",
                    "DISTINGUE_ANUAL_ACUMULADO": "NO_DIRECTO",
                }
            )
            granular_rows.append(
                {
                    "ID_FUENTE": row["ID_FUENTE"],
                    "RUTA": row["RUTA"],
                    "HOJA": row["HOJA"],
                    "GRANULARIDAD": row["GRANULARIDAD"],
                    "POSIBLE_USO": row["POSIBLE_USO"],
                }
            )
            # Solo se muestrean catálogos para fuentes chicas/clave, evitando exponer datos personales completos.
            if row["NOMBRE"] == "PROMEDIOSDEALUMNOS_7804.xlsx" and row["HOJA"] in {"Hoja1", "DatosAlumnos"}:
                df = load_source_dataframe(row["ID_FUENTE"], row["RUTA"], row["HOJA"]).head(5000)
                for col in headers:
                    if any(k in norm(col) for k in ["ANO", "PERIODO", "ESTADO", "CONVALIDADO", "NIVEL"]):
                        values = df[col].astype(str).str.strip().value_counts(dropna=False).head(40)
                        for value, count in values.items():
                            catalog_rows.append(
                                {
                                    "ID_FUENTE": row["ID_FUENTE"],
                                    "HOJA": row["HOJA"],
                                    "COLUMNA": col,
                                    "VALOR_OBSERVADO": value,
                                    "CASOS_MUESTRA": int(count),
                                }
                            )
                if "ANO" in df.columns:
                    temporal_rows[-1]["CONTIENE_2025"] = "SI" if df["ANO"].astype(str).eq("2025").any() else "NO_EN_MUESTRA"
                    temporal_rows[-1]["RANGO_TEMPORAL_OBSERVADO"] = join_unique(df["ANO"].astype(str), limit=30)
        dict_df = pd.DataFrame(dict_rows)
        cat_df = pd.DataFrame(catalog_rows)
        temp_df = pd.DataFrame(temporal_rows)
        gran_df = pd.DataFrame(granular_rows)
        self.profile = candidates
        write_xlsx(
            self.run_dir / "02_PERFIL_FUENTES" / "PERFIL_FUENTES_CANDIDATAS.xlsx",
            {
                "FUENTES_CANDIDATAS": candidates,
                "DICCIONARIO_COLUMNAS": dict_df,
                "CATALOGOS_OBSERVADOS": cat_df,
                "RANGOS_TEMPORALES": temp_df,
                "GRANULARIDAD": gran_df,
            },
        )
        write_tsv(self.run_dir / "02_PERFIL_FUENTES" / "DICCIONARIO_COLUMNAS_OBSERVADAS.tsv", dict_df)
        write_tsv(self.run_dir / "02_PERFIL_FUENTES" / "VALORES_CATALOGO_OBSERVADOS.tsv", cat_df)
        write_tsv(self.run_dir / "02_PERFIL_FUENTES" / "RANGOS_TEMPORALES.tsv", temp_df)
        write_tsv(self.run_dir / "02_PERFIL_FUENTES" / "ESTRUCTURA_GRANULARIDAD.tsv", gran_df)
        self.validation("A2", "PERFIL_CONTENIDO_GENERADO", "OK", f"candidatas={len(candidates)}")
        self.mark("A2", "COMPLETADA")
        self.checkpoint("A2", "COMPLETADA", [self.run_dir / "02_PERFIL_FUENTES" / "PERFIL_FUENTES_CANDIDATAS.xlsx"], "A3")

    def select_core_academic_sources(self) -> pd.DataFrame:
        profile = self.profile if not self.profile.empty else pd.read_csv(self.run_dir / "02_PERFIL_FUENTES" / "ESTRUCTURA_GRANULARIDAD.tsv", sep="\t", dtype=str, keep_default_na=False)
        inv = self.inventory if not self.inventory.empty else pd.read_csv(self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.tsv", sep="\t", dtype=str, keep_default_na=False)
        candidates = inv[inv["ESTADO_VALIDACION"].eq("CANDIDATA_POR_CONTENIDO")].copy()
        # Priorización por contenido, no por nombre como prueba: debe tener columnas detectadas.
        core = candidates[
            candidates["CATEGORIAS_DETECTADAS"].str.contains("CODCLI|RUT", na=False)
            & candidates["CATEGORIAS_DETECTADAS"].str.contains("PERIODO|ASIGNATURA|ESTADO|CONVALIDACION", na=False)
        ].copy()
        return core

    def a3(self) -> pd.DataFrame:
        self.mark("A3", "EN_EJECUCION")
        core = self.select_core_academic_sources()
        rows = []
        for field in SIES_FIELDS:
            if field == "VIGENCIA":
                rows.append(
                    {
                        "CAMPO_SIES": field,
                        "REGLA_OFICIAL": "Instructivo: mantener 1 para registros del universo; 0 solo elimina registro cargado por error.",
                        "FUENTE_CANDIDATA": "Instructivo oficial",
                        "COLUMNA_ORIGEN": "VIGENCIA",
                        "GRANULARIDAD": "REGLA_OFICIAL",
                        "PERIODO_CUBIERTO": "Proceso 2026 / datos 2025",
                        "LLAVE_DISPONIBLE": "No aplica",
                        "TRANSFORMACION_NECESARIA": "Ninguna en etapa preparatoria; no se materializa.",
                        "DEPENDENCIA_CARRERAS": "NO",
                        "INCLUYE_CONVALIDACIONES": "NO_APLICA",
                        "DISTINGUE_ANUAL_ACUMULADO": "NO_APLICA",
                        "NIVEL_RESPALDO": "OFICIAL",
                        "RIESGO": "BAJO",
                        "ESTADO": "FUENTE_DIRECTA_DEMOSTRADA",
                    }
                )
                continue
            if field == "PLAN_ESTUDIOS":
                rows.append(
                    {
                        "CAMPO_SIES": field,
                        "REGLA_OFICIAL": "Debe existir en Carreras para el CODIGO_UNICO del estudiante.",
                        "FUENTE_CANDIDATA": str(self.exec_dir / "03_CONCILIACION_CARRERAS" / "MAPEO_CODIGO_UNICO_PLAN.tsv"),
                        "COLUMNA_ORIGEN": "PLAN_ESTUDIOS / CODPESTUD_PLAN_INSTITUCIONAL",
                        "GRANULARIDAD": "CODIGO_UNICO + PLAN",
                        "PERIODO_CUBIERTO": "2025 preparatorio",
                        "LLAVE_DISPONIBLE": "CODIGO_UNICO + PLAN_ESTUDIOS",
                        "TRANSFORMACION_NECESARIA": "Validar dependencia Carreras; no materializar.",
                        "DEPENDENCIA_CARRERAS": "SI",
                        "INCLUYE_CONVALIDACIONES": "NO_APLICA",
                        "DISTINGUE_ANUAL_ACUMULADO": "NO_APLICA",
                        "NIVEL_RESPALDO": "CONCILIACION_INTERNA_TRAZADA",
                        "RIESGO": "ALTO_POR_CARRERAS_BLOQUEADAS",
                        "ESTADO": "FUENTE_PARCIAL",
                    }
                )
                continue
            if core.empty:
                rows.append(
                    {
                        "CAMPO_SIES": field,
                        "REGLA_OFICIAL": "Instructivo Anexo II",
                        "FUENTE_CANDIDATA": "",
                        "COLUMNA_ORIGEN": "",
                        "GRANULARIDAD": "",
                        "PERIODO_CUBIERTO": "",
                        "LLAVE_DISPONIBLE": "",
                        "TRANSFORMACION_NECESARIA": "",
                        "DEPENDENCIA_CARRERAS": "SI",
                        "INCLUYE_CONVALIDACIONES": "NO_DETERMINADO",
                        "DISTINGUE_ANUAL_ACUMULADO": "NO",
                        "NIVEL_RESPALDO": "SIN_FUENTE",
                        "RIESGO": "ALTO",
                        "ESTADO": "SIN_FUENTE",
                    }
                )
            else:
                best = core.iloc[0]
                rows.append(
                    {
                        "CAMPO_SIES": field,
                        "REGLA_OFICIAL": "Instructivo Anexo II; no incluir convalidaciones en anual; acumulado puede incluir reconocimientos permitidos.",
                        "FUENTE_CANDIDATA": f"{best['RUTA']}::{best['HOJA']}",
                        "COLUMNA_ORIGEN": best["ENCABEZADOS"],
                        "GRANULARIDAD": best["GRANULARIDAD"],
                        "PERIODO_CUBIERTO": "Contiene año/período en perfil; requiere filtro 2025 para anual.",
                        "LLAVE_DISPONIBLE": "RUT/CODCLI + carrera/plan según columnas detectadas",
                        "TRANSFORMACION_NECESARIA": "Agrupación por estudiante-carrera-plan-periodo; exclusión anual de convalidaciones; no ejecutada.",
                        "DEPENDENCIA_CARRERAS": "SI",
                        "INCLUYE_CONVALIDACIONES": "COLUMNA_OBSERVADA" if "CONVALIDACION" in str(best["CATEGORIAS_DETECTADAS"]) else "NO_DETERMINADO",
                        "DISTINGUE_ANUAL_ACUMULADO": "PARCIAL",
                        "NIVEL_RESPALDO": "FUENTE_INSTITUCIONAL_OBSERVADA",
                        "RIESGO": "ALTO_HASTA_VALIDAR_REGLA_OPERACIONAL",
                        "ESTADO": "FUENTE_PARCIAL",
                    }
                )
        matrix = pd.DataFrame(rows)
        write_tsv(self.run_dir / "03_MAPEO_CAMPOS" / "MATRIZ_FUENTE_CAMPO_SIES.tsv", matrix)
        write_xlsx(self.run_dir / "03_MAPEO_CAMPOS" / "MATRIZ_FUENTE_CAMPO_SIES.xlsx", {"MATRIZ": matrix})
        write_tsv(self.run_dir / "03_MAPEO_CAMPOS" / "CAMPOS_SIN_FUENTE.tsv", matrix[matrix["ESTADO"].eq("SIN_FUENTE")])
        write_tsv(self.run_dir / "03_MAPEO_CAMPOS" / "CAMPOS_CON_FUENTE_PARCIAL.tsv", matrix[matrix["ESTADO"].eq("FUENTE_PARCIAL")])
        self.validation("A3", "MATRIZ_FUENTE_CAMPO_GENERADA", "OK", matrix["ESTADO"].value_counts().to_dict().__repr__())
        self.mark("A3", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A3", "COMPLETADA_CON_BRECHAS", [self.run_dir / "03_MAPEO_CAMPOS" / "MATRIZ_FUENTE_CAMPO_SIES.xlsx"], "A4")
        return matrix

    def source_key_columns(self, df: pd.DataFrame) -> dict[str, str]:
        cols = {norm(c): c for c in df.columns}
        result = {}
        for key in ["RUT", "CODCLI", "CODCARR", "CODCARPR", "CODIGO_UNICO", "PLAN_DE_ESTUDIO", "PLAN_ESTUDIOS", "ANO", "PERIODO"]:
            nkey = norm(key)
            for ncol, col in cols.items():
                if nkey in ncol:
                    result[key] = col
                    break
        return result

    def coverage_for_source(self, row: pd.Series, df5809: pd.DataFrame, identity: pd.DataFrame) -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        try:
            src = load_source_dataframe(row["ID_FUENTE"], row["RUTA"], row["HOJA"])
        except Exception as exc:
            return (
                {
                    "ID_FUENTE": row["ID_FUENTE"],
                    "RUTA": row["RUTA"],
                    "HOJA": row["HOJA"],
                    "ERROR": f"{type(exc).__name__}: {exc}",
                },
                pd.DataFrame(),
                pd.DataFrame(),
                pd.DataFrame(),
            )
        keys = self.source_key_columns(src)
        src = src.copy()
        if "RUT" in keys:
            src["DOC_BODY_SRC"] = src[keys["RUT"]].map(lambda value: split_rut(value)[0])
        elif "NUM_DOCUMENTO" in keys:
            src["DOC_BODY_SRC"] = src[keys["NUM_DOCUMENTO"]].map(compact)
        else:
            src["DOC_BODY_SRC"] = ""
        if "CODCLI" in keys:
            src["CODCLI_SRC"] = src[keys["CODCLI"]].map(lambda value: str(value).strip())
        else:
            src["CODCLI_SRC"] = ""
        carrera_col = keys.get("CODCARR") or keys.get("CODCARPR") or keys.get("CODIGO_UNICO")
        src["CARRERA_SRC"] = src[carrera_col].map(lambda value: str(value).strip()) if carrera_col else ""
        src_doc_set = set(src.loc[src["DOC_BODY_SRC"].ne(""), "DOC_BODY_SRC"])
        src_codcli_set = set(src.loc[src["CODCLI_SRC"].ne(""), "CODCLI_SRC"])
        doc_matches = df5809[df5809["DOC_PRECARGA_BODY"].isin(src_doc_set)][["ID_FILA_5809", "DOC_PRECARGA_BODY", "CODIGO_UNICO"]].copy()
        ident_exploded = []
        for _, ir in identity.iterrows():
            for codcli in str(ir.get("CODCLI_LISTA", "")).split(" | "):
                codcli = codcli.strip()
                if codcli:
                    ident_exploded.append({"ID_FILA_5809": ir["ID_FILA_5809"], "CODCLI": codcli, "CLASIFICACION_IDENTIDAD": ir["CLASIFICACION_IDENTIDAD"]})
        ident_ex = pd.DataFrame(ident_exploded)
        codcli_matches = pd.DataFrame()
        if not ident_ex.empty and src_codcli_set:
            codcli_matches = ident_ex[ident_ex["CODCLI"].isin(src_codcli_set)].copy()
        multiple = pd.DataFrame()
        if "DOC_BODY_SRC" in src.columns:
            counts = src[src["DOC_BODY_SRC"].ne("")].groupby("DOC_BODY_SRC").size().reset_index(name="FILAS_FUENTE")
            multiple = counts[counts["FILAS_FUENTE"].gt(1)].head(10000)
        conflicts = pd.DataFrame()
        if not doc_matches.empty and carrera_col:
            expected = identity[["ID_FILA_5809", "CODCARR_ESPERADO_SEGUN_CARRERAS"]].merge(doc_matches, on="ID_FILA_5809", how="inner")
            carreras_src = src.groupby("DOC_BODY_SRC")["CARRERA_SRC"].agg(join_unique).reset_index()
            expected = expected.merge(carreras_src, left_on="DOC_PRECARGA_BODY", right_on="DOC_BODY_SRC", how="left")
            conflicts = expected[
                expected["CODCARR_ESPERADO_SEGUN_CARRERAS"].astype(str).ne("")
                & expected["CARRERA_SRC"].astype(str).ne("")
                & ~expected.apply(lambda r: str(r["CODCARR_ESPERADO_SEGUN_CARRERAS"]) in str(r["CARRERA_SRC"]).split(" | "), axis=1)
            ].copy()
        summary = {
            "ID_FUENTE": row["ID_FUENTE"],
            "RUTA": row["RUTA"],
            "HOJA": row["HOJA"],
            "FILAS_FUENTE": len(src),
            "PERSONAS_UNICAS_DOC_FUENTE": len(src_doc_set),
            "CODCLI_UNICOS_FUENTE": len(src_codcli_set),
            "COINCIDENCIAS_DOCUMENTO_EXACTO": int(doc_matches["ID_FILA_5809"].nunique()) if not doc_matches.empty else 0,
            "COINCIDENCIAS_CODCLI": int(codcli_matches["ID_FILA_5809"].nunique()) if not codcli_matches.empty else 0,
            "COINCIDENCIAS_DOCUMENTO_CARRERA": "",
            "COINCIDENCIAS_CODCLI_CARRERA": "",
            "SOLO_FUENTE_DOC": max(len(src_doc_set) - (doc_matches["DOC_PRECARGA_BODY"].nunique() if not doc_matches.empty else 0), 0),
            "SOLO_5809_DOC": max(df5809["DOC_PRECARGA_BODY"].nunique() - (doc_matches["DOC_PRECARGA_BODY"].nunique() if not doc_matches.empty else 0), 0),
            "MULTIPLES_COINCIDENCIAS": len(multiple),
            "CONFLICTOS_CARRERA": len(conflicts),
            "CONFLICTOS_DOCUMENTO": "",
        }
        return summary, multiple, conflicts, codcli_matches

    def a4(self) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        self.mark("A4", "EN_EJECUCION")
        if self.df5809.empty:
            self.df5809 = read_5809_from_prep(self.prep_dir)
        if self.identity.empty:
            self.identity = read_identity(self.prep_dir)
        core = self.select_core_academic_sources()
        # Mantiene el diagnóstico focalizado y auditable.
        core = core.head(25).copy()
        cov_rows = []
        multi_rows = []
        conflict_rows = []
        codcli_rows = []
        for _, row in core.iterrows():
            summary, multiple, conflicts, codcli = self.coverage_for_source(row, self.df5809, self.identity)
            cov_rows.append(summary)
            if not multiple.empty:
                multiple.insert(0, "ID_FUENTE", row["ID_FUENTE"])
                multiple.insert(1, "HOJA", row["HOJA"])
                multi_rows.append(multiple)
            if not conflicts.empty:
                conflicts.insert(0, "ID_FUENTE", row["ID_FUENTE"])
                conflicts.insert(1, "HOJA", row["HOJA"])
                conflict_rows.append(conflicts)
            if not codcli.empty:
                codcli.insert(0, "ID_FUENTE", row["ID_FUENTE"])
                codcli.insert(1, "HOJA", row["HOJA"])
                codcli_rows.append(codcli)
        cov = pd.DataFrame(cov_rows)
        multi = pd.concat(multi_rows, ignore_index=True) if multi_rows else pd.DataFrame()
        conflicts = pd.concat(conflict_rows, ignore_index=True) if conflict_rows else pd.DataFrame()
        codcli_matches = pd.concat(codcli_rows, ignore_index=True) if codcli_rows else pd.DataFrame()
        by_identity = []
        for _, row in core.iterrows():
            summary, _, _, codcli = self.coverage_for_source(row, self.df5809, self.identity)
            matched_ids = set()
            # Documento.
            try:
                src = load_source_dataframe(row["ID_FUENTE"], row["RUTA"], row["HOJA"])
                keys = self.source_key_columns(src)
                if "RUT" in keys:
                    doc_set = set(src[keys["RUT"]].map(lambda value: split_rut(value)[0]))
                    matched_ids.update(self.df5809[self.df5809["DOC_PRECARGA_BODY"].isin(doc_set)]["ID_FILA_5809"])
            except Exception:
                pass
            if not codcli.empty:
                matched_ids.update(codcli["ID_FILA_5809"])
            tmp = self.identity.copy()
            tmp["CUBIERTO"] = tmp["ID_FILA_5809"].isin(matched_ids)
            for cl, grp in tmp.groupby("CLASIFICACION_IDENTIDAD"):
                by_identity.append({"ID_FUENTE": row["ID_FUENTE"], "HOJA": row["HOJA"], "CLASIFICACION_IDENTIDAD": cl, "FILAS": len(grp), "CUBIERTAS": int(grp["CUBIERTO"].sum())})
        by_ident = pd.DataFrame(by_identity)
        no_coverage_ids = set(self.df5809["ID_FILA_5809"])
        if not cov.empty and cov["COINCIDENCIAS_DOCUMENTO_EXACTO"].max() > 0:
            best = cov.sort_values(["COINCIDENCIAS_DOCUMENTO_EXACTO", "COINCIDENCIAS_CODCLI"], ascending=False).iloc[0]
            src = load_source_dataframe(best["ID_FUENTE"], best["RUTA"], best["HOJA"])
            keys = self.source_key_columns(src)
            if "RUT" in keys:
                doc_set = set(src[keys["RUT"]].map(lambda value: split_rut(value)[0]))
                no_coverage_ids = set(self.df5809.loc[~self.df5809["DOC_PRECARGA_BODY"].isin(doc_set), "ID_FILA_5809"])
        no_cov = self.df5809[self.df5809["ID_FILA_5809"].isin(no_coverage_ids)][["ID_FILA_5809", "NUM_DOCUMENTO", "DV", "CODIGO_UNICO"]].copy()
        write_tsv(self.run_dir / "04_COBERTURA" / "COBERTURA_POR_FUENTE.tsv", cov)
        write_tsv(self.run_dir / "04_COBERTURA" / "COBERTURA_POR_TIPO_IDENTIDAD.tsv", by_ident)
        write_tsv(self.run_dir / "04_COBERTURA" / "COINCIDENCIAS_MULTIPLES.tsv", multi)
        write_tsv(self.run_dir / "04_COBERTURA" / "CONFLICTOS_CARRERA.tsv", conflicts)
        write_tsv(self.run_dir / "04_COBERTURA" / "CONFLICTOS_DOCUMENTO.tsv", pd.DataFrame(columns=["ID_FUENTE", "DETALLE"]))
        write_tsv(self.run_dir / "04_COBERTURA" / "CASOS_SIN_COBERTURA.tsv", no_cov)
        write_xlsx(
            self.run_dir / "04_COBERTURA" / "COBERTURA_FUENTES_VS_5809.xlsx",
            {
                "COBERTURA_FUENTE": cov,
                "COBERTURA_IDENTIDAD": by_ident,
                "MULTIPLES": multi.head(100000),
                "CONFLICTOS_CARRERA": conflicts.head(100000),
                "SIN_COBERTURA": no_cov,
            },
        )
        self.validation("A4", "COBERTURA_CALCULADA", "OK", f"fuentes evaluadas={len(core)}")
        self.mark("A4", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A4", "COMPLETADA_CON_BRECHAS", [self.run_dir / "04_COBERTURA" / "COBERTURA_FUENTES_VS_5809.xlsx"], "A5")
        return cov, by_ident, no_cov

    def a5(self) -> pd.DataFrame:
        self.mark("A5", "EN_EJECUCION")
        if self.identity.empty:
            self.identity = read_identity(self.prep_dir)
        review = self.identity[~self.identity["CLASIFICACION_IDENTIDAD"].eq("IDENTIDAD_EXACTA")].copy()
        rows = []
        for _, r in review.iterrows():
            classif = r["CLASIFICACION_IDENTIDAD"]
            if classif == "MULTIPLES_CODCLI":
                estado = "MULTIPLES_CODCLI_VARIAS_TRAYECTORIAS"
                if str(r.get("CODCARR_ESPERADO_SEGUN_CARRERAS", "")) and str(r.get("CODCARR_ESPERADO_SEGUN_CARRERAS", "")) in str(r.get("CODCARPR_LISTA", "")).split(" | "):
                    estado = "MULTIPLES_CODCLI_UNO_CON_TRAYECTORIA_COMPATIBLE"
            elif classif == "CONTRADICCION_CARRERA":
                estado = "CONTRADICCION_CARRERA"
            elif classif == "CONTRADICCION_DOCUMENTO":
                estado = "CONTRADICCION_DOCUMENTO"
            elif classif == "SIN_IDENTIDAD":
                estado = "SIN_IDENTIDAD"
            else:
                estado = "PENDIENTE_EVIDENCIA"
            rows.append(
                {
                    "ID_FILA_5809": r["ID_FILA_5809"],
                    "RUT_5809": f"{r['NUM_DOCUMENTO']}-{r['DV']}",
                    "RUT_INSTITUCIONAL": r.get("RUT_INSTITUCIONAL", ""),
                    "RUT_HISTORICO": r.get("RUT_INSTITUCIONAL", "") if classif == "CONTRADICCION_DOCUMENTO" else "",
                    "CODCLI_DISPONIBLES": r.get("CODCLI_LISTA", ""),
                    "CODCLI_CON_DATOS_2025": "PENDIENTE_CRUCE_FUENTE_AVANCE",
                    "CARRERA_5809": r.get("CODIGO_UNICO", ""),
                    "CARRERAS_HISTORICAS": r.get("CODCARPR_LISTA", ""),
                    "PLANES_HISTORICOS": "PENDIENTE",
                    "FUENTES_COINCIDENTES": "PROMEDIOSDEALUMNOS::DatosAlumnos" if classif != "SIN_IDENTIDAD" else "",
                    "ESTADO_IDENTIDAD": estado,
                    "EVIDENCIA": r.get("CLASIFICACION_IDENTIDAD", ""),
                    "CONTRADICCION": "SI" if "CONTRADICCION" in estado else "NO",
                    "DECISION_PROPUESTA": "USAR_SOLO_PARA_CONSULTA_CON_TRAZABILIDAD" if estado not in {"CONTRADICCION_DOCUMENTO", "SIN_IDENTIDAD"} else "NO_USAR_SIN_RESOLUCION",
                    "PUEDE_USARSE_PARA_CONSULTAR_AVANCE": "SI_CON_CONTROL" if estado not in {"CONTRADICCION_DOCUMENTO", "SIN_IDENTIDAD"} else "NO",
                    "PUEDE_USARSE_EN_CARGA": "NO",
                }
            )
        exp = pd.DataFrame(rows)
        write_tsv(self.run_dir / "05_IDENTIDAD" / "IDENTIDADES_RESUELTAS.tsv", exp[exp["PUEDE_USARSE_PARA_CONSULTAR_AVANCE"].eq("SI_CON_CONTROL")])
        write_tsv(self.run_dir / "05_IDENTIDAD" / "MULTIPLES_TRAYECTORIAS.tsv", exp[exp["ESTADO_IDENTIDAD"].str.contains("MULTIPLES", na=False)])
        write_tsv(self.run_dir / "05_IDENTIDAD" / "CONTRADICCIONES_IDENTIDAD.tsv", exp[exp["CONTRADICCION"].eq("SI")])
        write_tsv(self.run_dir / "05_IDENTIDAD" / "CASOS_SIN_IDENTIDAD.tsv", exp[exp["ESTADO_IDENTIDAD"].eq("SIN_IDENTIDAD")])
        write_xlsx(
            self.run_dir / "05_IDENTIDAD" / "EXPEDIENTE_IDENTIDAD_1008.xlsx",
            {
                "RESUMEN": exp["ESTADO_IDENTIDAD"].value_counts().rename_axis("ESTADO").reset_index(name="CASOS"),
                "EXPEDIENTE": exp,
                "CONTRADICCIONES": exp[exp["CONTRADICCION"].eq("SI")],
            },
        )
        self.validation("A5", "EXPEDIENTE_IDENTIDAD_1008", "OK", f"filas={len(exp)}")
        self.mark("A5", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A5", "COMPLETADA_CON_BRECHAS", [self.run_dir / "05_IDENTIDAD" / "EXPEDIENTE_IDENTIDAD_1008.xlsx"], "A6")
        return exp

    def a6(self) -> pd.DataFrame:
        self.mark("A6", "EN_EJECUCION")
        if self.df5809.empty:
            self.df5809 = read_5809_from_prep(self.prep_dir)
        if self.identity.empty:
            self.identity = read_identity(self.prep_dir)
        cov = pd.read_csv(self.run_dir / "04_COBERTURA" / "COBERTURA_POR_FUENTE.tsv", sep="\t", dtype=str, keep_default_na=False)
        best_has_academic = not cov.empty and pd.to_numeric(cov["COINCIDENCIAS_DOCUMENTO_EXACTO"], errors="coerce").fillna(0).max() > 0
        plan = pd.read_csv(self.prep_dir / "04_PLANES" / "PLAN_UNICO.tsv", sep="\t", dtype=str, keep_default_na=False)
        dep_ids = set(pd.read_csv(self.prep_dir / "04_PLANES" / "DEPENDENCIAS_CARRERAS.tsv", sep="\t", dtype=str, keep_default_na=False)["ID_FILA_5809"])
        identity_by_id = self.identity.set_index("ID_FILA_5809")["CLASIFICACION_IDENTIDAD"].to_dict()
        rows = []
        for _, r in self.df5809.iterrows():
            rid = r["ID_FILA_5809"]
            id_state = identity_by_id.get(rid, "SIN_IDENTIDAD")
            requires_identity = "NO" if id_state == "IDENTIDAD_EXACTA" else "SI"
            requires_carreras = "SI" if rid in dep_ids else "NO"
            if id_state in {"CONTRADICCION_DOCUMENTO", "SIN_IDENTIDAD"}:
                estado = "NO_PREPARABLE_POR_CONTRADICCION"
            elif not best_has_academic:
                estado = "NO_PREPARABLE_SIN_FUENTE_ACADEMICA"
            elif requires_carreras == "SI" and requires_identity == "SI":
                estado = "PREPARABLE_TRAS_AMBAS"
            elif requires_carreras == "SI":
                estado = "PREPARABLE_TRAS_CONFIRMACION_CARRERAS"
            elif requires_identity == "SI":
                estado = "PREPARABLE_TRAS_RESOLVER_IDENTIDAD"
            else:
                estado = "PREPARABLE_CON_FUENTES_ACTUALES"
            possible = "SI_CON_DEPENDENCIAS" if estado.startswith("PREPARABLE") else "NO"
            rows.append(
                {
                    "ID_FILA_5809": rid,
                    "PUEDE_DETERMINAR_CURSO_1ER_SEM": possible,
                    "PUEDE_DETERMINAR_CURSO_2DO_SEM": possible,
                    "PUEDE_DETERMINAR_UNIDADES_CURSADAS_2025": possible,
                    "PUEDE_DETERMINAR_UNIDADES_APROBADAS_2025": possible,
                    "PUEDE_DETERMINAR_CURSADAS_TOTAL": possible,
                    "PUEDE_DETERMINAR_APROBADAS_TOTAL": possible,
                    "PUEDE_DETERMINAR_PLAN": "SI_PREPARATORIO" if rid not in set(pd.read_csv(self.prep_dir / "04_PLANES" / "PLAN_NO_ENCONTRADO.tsv", sep="\t", dtype=str, keep_default_na=False).get("ID_FILA_5809", [])) else "NO",
                    "PUEDE_DETERMINAR_UNIDAD_MEDIDA": "NO_HASTA_CONFIRMAR_CARRERAS" if requires_carreras == "SI" else "SI_PREPARATORIO",
                    "FUENTE_PRINCIPAL": "PROMEDIOSDEALUMNOS::Hoja1" if best_has_academic else "",
                    "FUENTES_COMPLEMENTARIAS": "DatosAlumnos; conciliacion Carreras; expediente periodizaciones",
                    "REQUIERE_CARRERAS": requires_carreras,
                    "REQUIERE_IDENTIDAD": requires_identity,
                    "REQUIERE_PERIODIZACION": "SI" if rid in set(pd.read_csv(self.prep_dir / "05_PERIODIZACIONES" / "CASOS_PENDIENTES_27.tsv", sep="\t", dtype=str, keep_default_na=False)["ID_FILA_5809"]) else "NO",
                    "ESTADO_PREPARACION": estado,
                }
            )
        fact = pd.DataFrame(rows)
        resumen = fact["ESTADO_PREPARACION"].value_counts().rename_axis("ESTADO_PREPARACION").reset_index(name="CASOS")
        write_tsv(self.run_dir / "06_CONCILIACION" / "RESUMEN_FACTIBILIDAD.tsv", resumen)
        write_tsv(self.run_dir / "06_CONCILIACION" / "FILAS_PREPARABLES.tsv", fact[fact["ESTADO_PREPARACION"].str.startswith("PREPARABLE")])
        write_tsv(self.run_dir / "06_CONCILIACION" / "FILAS_NO_PREPARABLES.tsv", fact[~fact["ESTADO_PREPARACION"].str.startswith("PREPARABLE")])
        write_xlsx(self.run_dir / "06_CONCILIACION" / "FACTIBILIDAD_CALCULO_2371.xlsx", {"RESUMEN": resumen, "FACTIBILIDAD": fact})
        self.validation("A6", "FACTIBILIDAD_SIN_CALCULO", "OK", resumen.to_dict(orient="records").__repr__())
        self.mark("A6", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A6", "COMPLETADA_CON_BRECHAS", [self.run_dir / "06_CONCILIACION" / "FACTIBILIDAD_CALCULO_2371.xlsx"], "A7")
        return fact

    def a7(self) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        self.mark("A7", "EN_EJECUCION")
        core = self.select_core_academic_sources()
        cat_rows = []
        conv_cases = []
        rec_cases = []
        risks = []
        for _, row in core.head(10).iterrows():
            try:
                df = load_source_dataframe(row["ID_FUENTE"], row["RUTA"], row["HOJA"])
            except Exception:
                continue
            cols = [c for c in df.columns if "CONVALID" in norm(c) or "RECONOC" in norm(c) or "HOMOLOG" in norm(c) or "ESTADO" in norm(c)]
            for col in cols:
                values = df[col].astype(str).str.strip().value_counts(dropna=False).head(100)
                for value, count in values.items():
                    cat_rows.append({"ID_FUENTE": row["ID_FUENTE"], "HOJA": row["HOJA"], "COLUMNA": col, "VALOR": value, "CASOS": int(count)})
                if "CONVALID" in norm(col):
                    cases = df[df[col].astype(str).str.strip().ne("")].head(5000).copy()
                    if not cases.empty:
                        cases.insert(0, "ID_FUENTE", row["ID_FUENTE"])
                        cases.insert(1, "HOJA", row["HOJA"])
                        conv_cases.append(cases)
            if not cols:
                risks.append({"ID_FUENTE": row["ID_FUENTE"], "HOJA": row["HOJA"], "RIESGO": "No se observaron columnas para separar convalidaciones/reconocimientos."})
        cat = pd.DataFrame(cat_rows)
        conv = pd.concat(conv_cases, ignore_index=True) if conv_cases else pd.DataFrame()
        rec = pd.concat(rec_cases, ignore_index=True) if rec_cases else pd.DataFrame()
        risk = pd.DataFrame(risks)
        pending_map = cat.copy()
        if not pending_map.empty:
            pending_map["MAPEO_OFICIAL"] = "PENDIENTE_VALIDACION_FUNCIONAL"
        write_tsv(self.run_dir / "07_BRECHAS" / "CATALOGO_OBSERVADO_ESTADOS_ASIGNATURA.tsv", cat)
        write_tsv(self.run_dir / "07_BRECHAS" / "MAPEO_PENDIENTE_ESTADOS_OFICIALES.tsv", pending_map)
        write_tsv(self.run_dir / "07_BRECHAS" / "CASOS_CON_CONVALIDACIONES.tsv", conv)
        write_tsv(self.run_dir / "07_BRECHAS" / "CASOS_CON_RECONOCIMIENTOS.tsv", rec)
        write_tsv(self.run_dir / "07_BRECHAS" / "RIESGOS_AVANCE_ANUAL_VS_ACUMULADO.tsv", risk)
        if cat.empty:
            self.block("A7", "CONVALIDACIONES_NO_SEPARABLES", "No se observaron columnas de estado/convalidacion en fuentes candidatas.", "Solicitar fuente academica con indicador de convalidacion/reconocimiento.")
        self.validation("A7", "CONVALIDACIONES_INSPECCIONADAS", "OK" if not cat.empty else "BRECHA", f"catalogo={len(cat)} casos_convalidacion={len(conv)}")
        self.mark("A7", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A7", "COMPLETADA_CON_BRECHAS", [self.run_dir / "07_BRECHAS" / "CATALOGO_OBSERVADO_ESTADOS_ASIGNATURA.tsv"], "A8")
        return cat, conv, risk

    def a8(self) -> pd.DataFrame:
        self.mark("A8", "EN_EJECUCION")
        identity_counts = read_identity(self.prep_dir)["CLASIFICACION_IDENTIDAD"].value_counts().to_dict()
        rows = []
        def add(tipo: str, universo: int, campo: str, disponible: str, faltante: str, accion: str, prioridad: str, impacto: str) -> None:
            rows.append(
                {
                    "ID_BRECHA": f"BR{len(rows)+1:04d}",
                    "TIPO": tipo,
                    "UNIVERSO_AFECTADO": universo,
                    "FILAS_AFECTADAS": "",
                    "CAMPO_SIES": campo,
                    "FUENTE_DISPONIBLE": disponible,
                    "FUENTE_FALTANTE": faltante,
                    "RESPONSABLE_SUGERIDO": "Institucion / equipo academico funcional",
                    "ACCION_REQUERIDA": accion,
                    "PRIORIDAD": prioridad,
                    "IMPACTO": impacto,
                    "ESTADO": "PENDIENTE",
                }
            )
        add("CARRERAS", 15, "PLAN_ESTUDIOS;UNIDADES_*", "Expediente 42/15 y solicitud periodizacion existente", "Semantica NIVEL; mapa NIVEL->anio; unidad de medida; unidades por anio", "Responder solicitud de periodizacion 15 planes.", "ALTA", "Bloquea Carreras final e integridad Matricula.")
        add("FUENTE_AVANCE", 2371, "CURSO_1ER_SEM;CURSO_2DO_SEM;UNIDADES_*", "PROMEDIOSDEALUMNOS::Hoja1 como fuente parcial", "Base academica validada con presencia semestral, cursadas/aprobadas anual y acumulado, con convalidaciones separables", "Solicitar fuente de avance 2025 con criterio de corte.", "ALTA", "Bloquea calculo de campos obligatorios.")
        add("IDENTIDAD", 1008, "IDENTIDAD_CODCLI", "Diagnostico identidad preparatorio", f"Múltiples/contradicciones/sin identidad: {identity_counts}", "Resolver casos requeridos por expediente identidad.", "MEDIA", "Puede impedir consulta de avance trazable.")
        add("PERIODIZACION", 27, "ANIO_CURRICULAR_DERIVADO_EVIDENCIA", "2 casos modelo directos", "Reporte individual o tabla institucional para 27 pendientes", "Solicitar reportes periodizacion 27.", "MEDIA", "Impide conversiones especiales sin supuestos.")
        bdf = pd.DataFrame(rows)
        write_tsv(self.run_dir / "07_BRECHAS" / "BRECHAS_CARRERAS.tsv", bdf[bdf["TIPO"].eq("CARRERAS")])
        write_tsv(self.run_dir / "07_BRECHAS" / "BRECHAS_FUENTE_AVANCE.tsv", bdf[bdf["TIPO"].eq("FUENTE_AVANCE")])
        write_tsv(self.run_dir / "07_BRECHAS" / "BRECHAS_IDENTIDAD.tsv", bdf[bdf["TIPO"].eq("IDENTIDAD")])
        write_tsv(self.run_dir / "07_BRECHAS" / "BRECHAS_PERIODIZACION.tsv", bdf[bdf["TIPO"].eq("PERIODIZACION")])
        write_xlsx(self.run_dir / "07_BRECHAS" / "EXPEDIENTE_BRECHAS_COMPLETO.xlsx", {"BRECHAS": bdf})
        self.validation("A8", "BRECHAS_CONSOLIDADAS", "OK", str(len(bdf)))
        self.mark("A8", "COMPLETADA_CON_BRECHAS")
        self.checkpoint("A8", "COMPLETADA_CON_BRECHAS", [self.run_dir / "07_BRECHAS" / "EXPEDIENTE_BRECHAS_COMPLETO.xlsx"], "A9")
        return bdf

    def a9(self) -> None:
        self.mark("A9", "EN_EJECUCION")
        period_req = self.exec_dir / "03_CONCILIACION_CARRERAS" / "EXPEDIENTE_42_PERIODIZACIONES_20260703_134419" / "SOLICITUD_PERIODIZACION_15_PLANES_20260703_134552" / "04_ENTREGA_INSTITUCIONAL" / "SOLICITUD_CONFIRMACION_PERIODIZACION_15_PLANES.xlsx"
        period_ref = pd.DataFrame(
            [
                {
                    "DOCUMENTO_EXISTENTE": str(period_req),
                    "EXISTE": period_req.exists(),
                    "SHA256": sha256(period_req) if period_req.exists() else "",
                    "ACCION": "Referenciar; no reemplazar.",
                }
            ]
        )
        fuente_cols = [
            "identificador_estudiante",
            "rut_o_documento",
            "codcli",
            "codigo_carrera",
            "plan",
            "anio",
            "periodo_academico",
            "presencia_o_inscripcion",
            "codigo_asignatura_unidad",
            "unidad_medida",
            "unidades_cursadas",
            "unidades_aprobadas",
            "estado_academico_unidad",
            "indicador_convalidacion_reconocimiento",
            "acumulados_si_existen",
            "fecha_extraccion",
            "criterio_corte",
        ]
        fuente_req = pd.DataFrame({"CAMPO_SOLICITADO": fuente_cols, "OBLIGATORIEDAD": "OBLIGATORIO", "OBSERVACION": "Requerido para preparar carga SIES; no modifica precarga."})
        ident = pd.read_excel(self.run_dir / "05_IDENTIDAD" / "EXPEDIENTE_IDENTIDAD_1008.xlsx", sheet_name="EXPEDIENTE", dtype=str, keep_default_na=False)
        ident_req = ident[ident["PUEDE_USARSE_PARA_CONSULTAR_AVANCE"].eq("NO") | ident["CONTRADICCION"].eq("SI")].copy()
        period27 = pd.read_csv(self.prep_dir / "05_PERIODIZACIONES" / "CASOS_PENDIENTES_27.tsv", sep="\t", dtype=str, keep_default_na=False)
        write_xlsx(self.run_dir / "08_SOLICITUDES" / "SOLICITUD_FUENTE_AVANCE_2025.xlsx", {"CAMPOS_REQUERIDOS": fuente_req, "REFERENCIA_PERIODIZACION_15": period_ref})
        write_xlsx(self.run_dir / "08_SOLICITUDES" / "SOLICITUD_RESOLUCION_IDENTIDAD.xlsx", {"CASOS_IDENTIDAD": ident_req})
        write_xlsx(self.run_dir / "08_SOLICITUDES" / "SOLICITUD_REPORTES_PERIODIZACION_27.xlsx", {"CASOS_27": period27})
        correo_avance = """# Solicitud fuente de avance curricular 2025

Proceso: Avance Curricular SIES 2026.
Año de referencia: 2025.
Universo: 2.371 filas de la precarga Matrícula Avance Curricular 2026.

Se solicita una base institucional para preparar la carga SIES. No se solicita modificar la precarga.

Campos mínimos: identificador de estudiante, RUT/documento, CODCLI, código de carrera, plan, año, período académico, presencia o inscripción, código de asignatura/unidad, unidad de medida, unidades cursadas, unidades aprobadas, estado académico de la unidad, indicador de convalidación/reconocimiento, acumulados si existen, fecha de extracción y criterio de corte.

Motivo e impacto: sin esta fuente no corresponde completar presencia, avance anual ni acumulados sin supuestos.
"""
        correo_ident = """# Solicitud resolución de identidad

Proceso: Avance Curricular SIES 2026.
Año de referencia: 2025.

Se adjunta expediente con casos que requieren intervención institucional por múltiples CODCLI, contradicción de documento, contradicción de carrera o ausencia de identidad. La información se usará solo para consultar avance y preparar la carga SIES; no autoriza modificar datos personales de la precarga.
"""
        correo_period = """# Solicitud reportes de periodización

Proceso: Avance Curricular SIES 2026.
Año de referencia: 2025.

Se solicita reporte individual de Avance de Malla o tabla institucional de periodización para los 27 casos pendientes y confirmación de los 15 planes ya solicitados. La planilla existente de 15 planes se referencia y no se reemplaza.

Impacto: sin esta evidencia no corresponde convertir niveles ni completar campos por semejanza.
"""
        (self.run_dir / "08_SOLICITUDES" / "CORREO_SOLICITUD_FUENTE_AVANCE.md").write_text(correo_avance, encoding="utf-8")
        (self.run_dir / "08_SOLICITUDES" / "CORREO_SOLICITUD_IDENTIDAD.md").write_text(correo_ident, encoding="utf-8")
        (self.run_dir / "08_SOLICITUDES" / "CORREO_SOLICITUD_PERIODIZACION.md").write_text(correo_period, encoding="utf-8")
        self.validation("A9", "PAQUETES_SOLICITUD_GENERADOS", "OK", "Fuente avance, identidad, periodizacion 27; referencia 15 planes verificada.")
        self.mark("A9", "COMPLETADA")
        self.checkpoint("A9", "COMPLETADA", [self.run_dir / "08_SOLICITUDES" / "SOLICITUD_FUENTE_AVANCE_2025.xlsx"], "A10")

    def a10(self) -> None:
        self.mark("A10", "EN_EJECUCION")
        outputs = []
        for path in sorted(self.run_dir.rglob("*")):
            if path.is_file() and path.name != "MANIFEST_SALIDAS.json":
                outputs.append({"RUTA": str(path), "SHA256": sha256(path), "TAMANO_BYTES": path.stat().st_size})
        write_json(self.paths["manifest_salidas"], {"fecha": now_iso(), "salidas": outputs})
        inv = pd.read_csv(self.run_dir / "01_INVENTARIO" / "INVENTARIO_DIRIGIDO_FUENTES_AVANCE.tsv", sep="\t", dtype=str, keep_default_na=False)
        matrix = pd.read_csv(self.run_dir / "03_MAPEO_CAMPOS" / "MATRIZ_FUENTE_CAMPO_SIES.tsv", sep="\t", dtype=str, keep_default_na=False)
        cov = pd.read_csv(self.run_dir / "04_COBERTURA" / "COBERTURA_POR_FUENTE.tsv", sep="\t", dtype=str, keep_default_na=False)
        fact = pd.read_csv(self.run_dir / "06_CONCILIACION" / "RESUMEN_FACTIBILIDAD.tsv", sep="\t", dtype=str, keep_default_na=False)
        brechas = pd.read_excel(self.run_dir / "07_BRECHAS" / "EXPEDIENTE_BRECHAS_COMPLETO.xlsx", sheet_name="BRECHAS", dtype=str)
        summary = pd.DataFrame(
            [
                ["ESTADO_FINAL", "DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS"],
                ["FUENTES_INVENTARIADAS_TABLAS_HOJAS", len(inv)],
                ["FUENTES_CANDIDATAS_UTILES", int(inv["ESTADO_VALIDACION"].eq("CANDIDATA_POR_CONTENIDO").sum())],
                ["CAMPOS_FUENTE_DIRECTA", int(matrix["ESTADO"].eq("FUENTE_DIRECTA_DEMOSTRADA").sum())],
                ["CAMPOS_FUENTE_PARCIAL", int(matrix["ESTADO"].eq("FUENTE_PARCIAL").sum())],
                ["CAMPOS_SIN_FUENTE", int(matrix["ESTADO"].eq("SIN_FUENTE").sum())],
                ["FUENTES_COBERTURA_EVALUADAS", len(cov)],
                ["BRECHAS", len(brechas)],
            ],
            columns=["INDICADOR", "VALOR"],
        )
        report = [
            "# Diagnostico de fuentes de avance",
            "",
            "Estado final: DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS",
            "",
            f"Carpeta: `{self.run_dir}`",
            "",
            f"Fuentes inventariadas (tablas/hojas): {len(inv)}.",
            f"Fuentes candidatas utiles: {int(inv['ESTADO_VALIDACION'].eq('CANDIDATA_POR_CONTENIDO').sum())}.",
            f"Campos con fuente directa: {int(matrix['ESTADO'].eq('FUENTE_DIRECTA_DEMOSTRADA').sum())}.",
            f"Campos con fuente parcial: {int(matrix['ESTADO'].eq('FUENTE_PARCIAL').sum())}.",
            f"Campos sin fuente: {int(matrix['ESTADO'].eq('SIN_FUENTE').sum())}.",
            "",
            "No se calcularon valores, no se genero CSV, no se marco SIES_READY y no se reanudo Carreras final.",
        ]
        (self.run_dir / "10_REPORTES" / "INFORME_DIAGNOSTICO_FUENTES_AVANCE.md").write_text("\n".join(report) + "\n", encoding="utf-8")
        write_xlsx(
            self.run_dir / "10_REPORTES" / "INFORME_DIAGNOSTICO_FUENTES_AVANCE.xlsx",
            {"RESUMEN": summary, "FACTIBILIDAD": fact, "BRECHAS": brechas},
        )
        # Actualizacion no destructiva del estado general.
        estado_path = self.exec_dir / "00_CONTROL" / "ESTADO_EJECUCION.json"
        estado = json.loads(estado_path.read_text(encoding="utf-8"))
        estado["diagnostico_fuentes_avance"] = {
            "estado": "DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS",
            "fecha": now_iso(),
            "ruta": str(self.run_dir),
            "csv_generado": False,
            "sies_ready_generado": False,
        }
        estado_path.write_text(json.dumps(estado, indent=2, ensure_ascii=False), encoding="utf-8")
        checkpoint_path = self.exec_dir / "00_CONTROL" / "CHECKPOINT_ACTUAL.json"
        checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        checkpoint["diagnostico_fuentes_avance"] = {
            "estado": "DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS",
            "ruta": str(self.run_dir),
            "informe": str(self.run_dir / "10_REPORTES" / "INFORME_DIAGNOSTICO_FUENTES_AVANCE.md"),
        }
        checkpoint_path.write_text(json.dumps(checkpoint, indent=2, ensure_ascii=False), encoding="utf-8")
        append_tsv(
            self.exec_dir / "00_CONTROL" / "VALIDACIONES.tsv",
            {"FASE": "A10", "VALIDACION": "DIAGNOSTICO_FUENTES_AVANCE", "ESTADO": "COMPLETADO_CON_BRECHAS", "DETALLE": str(self.run_dir)},
            ["FASE", "VALIDACION", "ESTADO", "DETALLE"],
        )
        append_tsv(
            self.exec_dir / "00_CONTROL" / "BLOQUEOS.tsv",
            {
                "FECHA": now_iso(),
                "FASE": "A10",
                "MOTIVO": "BRECHAS_FUENTE_AVANCE_IDENTIFICADAS",
                "DETALLE": "Fuente academica parcial; se requieren confirmaciones y solicitud formal.",
                "ACCION": "Revisar paquetes de solicitud A9.",
            },
            ["FECHA", "FASE", "MOTIVO", "DETALLE", "ACCION"],
        )
        (self.exec_dir / "00_CONTROL" / "REANUDAR_DESDE_AQUI.md").write_text(
            "\n".join(
                [
                    "# Reanudar cierre integral",
                    "",
                    "Estado: Diagnostico de fuentes de avance completado con brechas.",
                    f"Ruta diagnostico: {self.run_dir}",
                    "",
                    "Siguiente accion: gestionar solicitudes A9 y resolver dependencias de Carreras/fuente de avance antes de generar candidatos.",
                    "",
                ]
            ),
            encoding="utf-8",
        )
        (self.exec_dir / "11_REPORTES").mkdir(exist_ok=True)
        shutil.copy2(self.run_dir / "10_REPORTES" / "INFORME_DIAGNOSTICO_FUENTES_AVANCE.md", self.exec_dir / "11_REPORTES" / "INFORME_CIERRE_PARCIAL_DIAGNOSTICO_FUENTES_AVANCE.md")
        self.mark("A10", "DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS")
        self.checkpoint("A10", "DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS", [self.run_dir / "10_REPORTES" / "INFORME_DIAGNOSTICO_FUENTES_AVANCE.md"], "A10")
        outputs = []
        for path in sorted(self.run_dir.rglob("*")):
            if path.is_file() and path.name != "MANIFEST_SALIDAS.json":
                outputs.append({"RUTA": str(path), "SHA256": sha256(path), "TAMANO_BYTES": path.stat().st_size})
        write_json(self.paths["manifest_salidas"], {"fecha": now_iso(), "salidas": outputs})

    def run(self) -> None:
        self.a0()
        self.a1()
        self.a2()
        self.a3()
        self.a4()
        self.a5()
        self.a6()
        self.a7()
        self.a8()
        self.a9()
        self.a10()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--execution-dir", type=Path, default=EXEC_DEFAULT)
    parser.add_argument("--prep-dir", type=Path, default=PREP_DEFAULT)
    parser.add_argument("--run-dir", type=Path)
    args = parser.parse_args()
    run_dir = init_run(args.execution_dir.resolve(), args.run_dir)
    runner = Runner(args.execution_dir.resolve(), args.prep_dir.resolve(), run_dir)
    runner.run()
    print("DIAGNOSTICO_FUENTES_COMPLETADO_CON_BRECHAS")
    print(f"Carpeta: {run_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
