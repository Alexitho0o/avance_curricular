#!/usr/bin/env python3
"""Reconstruye oferta publicada y su cobertura MRUN, sin resolver mallas UMAS.

Contrato: docs/fase_3_oferta_historica_sies.md. Solo lectura de XLSX (openpyxl).
Exige descriptor explícito con rutas/huellas y auditoría de extracción original.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import unicodedata

CODE = re.compile(r'I(\d+)S(\d+)C(\d+)J(\d+)V(\d+)\Z')
JORNADA = {'1':'Diurna', '2':'Vespertina', '3':'Semipresencial', '4':'A Distancia'}
MODALIDAD = {'1':'Presencial', '2':'Semipresencial', '3':'No Presencial'}
PLAN = {'1':'Plan Regular', '3':'Plan Regular de Continuidad'}
VIGENCIA = {'1':'Vigente con estudiantes nuevos', '2':'Vigente sin estudiantes nuevos'}
# Esquema de la publicación consolidada, NO el instructivo de carga CFT/IP 2026.
NIVEL_GLOBAL = {'1':'Pregrado', '3':'Postítulo'}
NIVEL_CARRERA = {'1':'Técnico de Nivel Superior', '2':'Profesional Sin Licenciatura',
                 '5':'Diplomado (desde un semestre)'}
CURRENT_FIELDS = {
    'CODIGO_IES_NUM':('CODIGO_IES', None), 'CODIGO_UNICO':('CODIGO_UNICO', None),
    'CODIGO_CARRERA':('CODIGO_CARRERA', None), 'NOMBRE_CARRERA':('NOMBRE_CARRERA', None),
    'MODALIDAD':('MODALIDAD', MODALIDAD), 'CODIGO_JORNADA':('JORNADA', JORNADA),
    'VERSION_CARRERA':('VERSION', None), 'CODIGO_TIPO_CARRERA':('TIPO_CARRERA', PLAN),
    'CARACT_PLAN_ESPECIAL':('PLAN_ESPECIAL', None),
    'DURACION_ESTUDIO':('DURACION_ESTUDIOS', None),
    'DURACION_TITULACION':('DURACION_TITULACION', None),
    'DURACION_TOTAL':('DURACION_TOTAL', None),
    'CODIGO_NIVEL_GLOBAL':('NIVEL_GLOBAL', NIVEL_GLOBAL),
    'CODIGO_NIVEL_CARRERA':('NIVEL_CARRERA', NIVEL_CARRERA), 'VIGENCIA':('VIGENCIA', VIGENCIA),
}
MRUN_FIELDS = {
    'CODIGO_IES':'cod_inst', 'CODIGO_SEDE':'cod_sede', 'CODIGO_CARRERA':'cod_carrera',
    'NOMBRE_CARRERA':'nomb_carrera', 'MODALIDAD':'modalidad', 'JORNADA':'jornada',
    'VERSION':'version', 'TIPO_CARRERA':'tipo_plan_carr', 'VIGENCIA':'vigencia_carrera',
    'DURACION_ESTUDIOS':'dur_estudio_carr', 'DURACION_TITULACION':'dur_proceso_tit',
    'DURACION_TOTAL':'dur_total_carr',
}
EXAMPLE_FAMILIES = ('I162S2C90J2','I162S2C77J4','I162S2C46J2',
                    'I162S2C46J4','I162S2C1J1','I162S2C35J4')


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True)+'\n').encode()


def text(value):
    if value is None:
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)


def normalized(value):
    """Comparación categórica sin cambiar valores conservados en las salidas."""
    value=''.join(c for c in unicodedata.normalize('NFD', value.upper())
                  if not unicodedata.combining(c))
    return ' '.join(value.split())


def compare(a, b, field):
    if not a.strip() or not b.strip():
        return 'SIN_DATO_AMBAS' if not a.strip() and not b.strip() else (
            'SIN_DATO_OFERTA' if not a.strip() else 'SIN_DATO_MRUN')
    if a==b:
        return 'IGUAL_LITERAL'
    na,nb=normalized(a),normalized(b)
    if field=='VIGENCIA':
        na,nb=na.replace('ALUMNOS','ESTUDIANTES'),nb.replace('ALUMNOS','ESTUDIANTES')
    return 'IGUAL_NORMALIZADO' if na==nb else 'DIFERENTE'


def family(code):
    p=CODE.fullmatch(code)
    return code.rsplit('V',1)[0] if p else None


def year(row):
    if not re.fullmatch(r'OFE_\d{4}', row['ANO']):
        raise ValueError('Año de oferta no reconocido: '+row['ANO'])
    return int(row['ANO'][4:])


def unique_index(rows, key, label):
    result={}
    for r in rows:
        k=key(r)
        if k in result:
            raise ValueError(f'Clave duplicada {label}: {k}')
        result[k]=r
    return result


def locator(row):
    return {'hoja':row['hoja_origen'], 'fila':int(row['fila_origen'])}


def profile(row):
    return {'anio':year(row), 'codigo_unico':row['CODIGO_UNICO'],
            **{f:row[f] for f in ('NOMBRE_CARRERA','MODALIDAD','JORNADA','VERSION',
               'TIPO_CARRERA','PLAN_ESPECIAL','DURACION_ESTUDIOS','DURACION_TITULACION',
               'DURACION_TOTAL','ANO_INICIO','VIGENCIA','NIVEL_GLOBAL','NIVEL_CARRERA')},
            'localizador':locator(row)}


def build(historic, mrun, current, coverage_years, institution='162', current_year=2026):
    """Construcción pura: fechas observadas no se convierten en vigencia de plan."""
    historic=sorted(historic, key=lambda r:(year(r),r['CODIGO_UNICO']))
    oi=unique_index(historic, lambda r:(year(r),r['CODIGO_UNICO']), 'oferta anual')
    mi=unique_index(mrun, lambda r:(int(r['cat_periodo']),r['codigo_unico']), 'MRUN anual')
    ci=unique_index(current, lambda r:r['CODIGO_UNICO'], 'oferta actual')
    if {year(r) for r in historic} != set(coverage_years):
        raise ValueError('Cobertura anual distinta del contrato')
    for r in historic:
        p=CODE.fullmatch(r['CODIGO_UNICO'])
        if not p or r['CODIGO_IES']!=institution or p[1]!=institution:
            raise ValueError('Código fuera de institución o esquema: '+r['CODIGO_UNICO'])
        if (p[1],p[2],p[3],p[5])!=tuple(r[f] for f in ('CODIGO_IES','CODIGO_SEDE','CODIGO_CARRERA','VERSION')):
            raise ValueError('Componentes incompatibles: '+r['CODIGO_UNICO'])
        if normalized(JORNADA[p[4]])!=normalized(r['JORNADA']):
            raise ValueError('Jornada incompatible: '+r['CODIGO_UNICO'])
    by_code=defaultdict(list); by_family=defaultdict(list); mr_code=defaultdict(list)
    for r in historic:
        by_code[r['CODIGO_UNICO']].append(r)
        by_family[family(r['CODIGO_UNICO'])].append(r)
    for r in mrun:
        if r['cod_inst']!=institution:
            raise ValueError('MRUN de otra institución')
        mr_code[r['codigo_unico']].append(r)

    changes=[]; histories=[]; start_relations=[]
    excluded={'ANO','CODIGO_UNICO','fila_origen','hoja_origen'}
    for code,rows in sorted(by_code.items()):
        years=sorted(year(r) for r in rows)
        for a,b in zip(rows,rows[1:]):
            changed={f:{'anterior':a[f],'posterior':b[f]} for f in sorted(a)
                     if f not in excluded and a[f]!=b[f]}
            if changed:
                changes.append({'codigo_unico':code,'anio_anterior':year(a),'anio_posterior':year(b),
                                'anios_consecutivos':year(b)==year(a)+1,
                                'localizador_anterior':locator(a),'localizador_posterior':locator(b),
                                'cambios':changed})
        for r in rows:
            start=r['ANO_INICIO']
            status='SIN_DATO' if not start else 'VALOR_CERO_SIN_ANIO' if start=='0' else 'FORMATO_NO_ANUAL'
            if re.fullmatch(r'\d{4}',start):
                status=('ANTERIOR' if int(start)<years[0] else
                        'POSTERIOR' if int(start)>years[0] else 'IGUAL')
            start_relations.append({'codigo_unico':code,'anio_observado':year(r),
                                    'primer_anio_oferta_observada':years[0],
                                    'ano_inicio_declarado':start,'relacion':status,
                                    'localizador':locator(r)})
        histories.append({'codigo_unico':code,'familia':family(code), 'version_informada':rows[0]['VERSION'],
                          'anios_oferta_observada':years, 'primer_anio_observado':years[0],
                          'ultimo_anio_observado':years[-1],
                          'anios_intermedios_sin_oferta_observada':sorted(set(range(years[0],years[-1]+1))-set(years)),
                          'nombre_literal_oferta_2026':ci.get(code,{}).get('NOMBRE_CARRERA'),
                          'historia':[profile(r) for r in rows]})
    families=[]; coexist=[]
    for fam,rows in sorted(by_family.items()):
        fy=defaultdict(list)
        for r in rows:fy[year(r)].append(r)
        versions=sorted({r['VERSION'] for r in rows}, key=int)
        first=min(fy)
        families.append({'familia':fam,'versiones_informadas':versions,
                         'primer_anio_observado':first,
                         'versiones_en_primer_anio':sorted({r['VERSION'] for r in fy[first]},key=int),
                         'codigos':sorted({r['CODIGO_UNICO'] for r in rows}),
                         'versiones_por_anio':{str(y):sorted({r['VERSION'] for r in rs},key=int) for y,rs in sorted(fy.items())}})
        for y,rs in sorted(fy.items()):
            if len(rs)>1:
                coexist.append({'familia':fam,'anio':y,'variantes':[profile(r) for r in rs]})

    coverage=[]
    for code in sorted(set(by_code)|set(mr_code)):
        os=by_code.get(code,[]); ms=mr_code.get(code,[])
        status='OFERTA_Y_MRUN' if os and ms else 'SOLO_OFERTA' if os else 'SOLO_MRUN'
        coverage.append({'codigo_unico':code,'familia':family(code), 'presencia':status,
                         'niveles_globales_oferta':' | '.join(sorted({r['NIVEL_GLOBAL'] for r in os})),
                         'anios_oferta':','.join(str(year(r)) for r in os),
                         'anios_mrun':','.join(str(v) for v in sorted({int(r['cat_periodo']) for r in ms})),
                         'oferta_2026':'SI' if code in ci else 'NO_OBSERVADA',
                         'nombre_literal_oferta_2026':ci.get(code,{}).get('NOMBRE_CARRERA',''),
                         'nombre_primera_oferta_observada':os[0]['NOMBRE_CARRERA'] if os else '',
                         'nombre_primera_matricula_observada':min(ms,key=lambda r:int(r['cat_periodo']))['nomb_carrera'] if ms else '',
                         'filas_matricula_mrun':sum(int(r['filas_matricula']) for r in ms)})
    annual=[]; comparisons=[]; comparable=Counter()
    for y,code in sorted(set(oi)|set(mi)):
        o=oi.get((y,code)); m=mi.get((y,code)); counts=Counter()
        status=('AMBAS_FUENTES' if o and m else 'SOLO_OFERTA' if o else
                'MRUN_FUERA_COBERTURA_OFERTA' if y not in coverage_years else 'MRUN_SIN_OFERTA_DEL_ANIO')
        if o and m:
            for of,mf in MRUN_FIELDS.items():
                state=compare(o[of],m[mf],of); counts[state]+=1; comparable[state]+=1
                comparisons.append({'anio':y,'codigo_unico':code,'campo_oferta':of,'campo_mrun':mf,
                                    'valor_oferta':o[of],'valor_mrun':m[mf],'estado':state,
                                    'fila_oferta':o['fila_origen'],'hoja_oferta':o['hoja_origen'],
                                    'registro_mrun_ejemplo':m['registro_ejemplo']})
        annual.append({'anio':y,'codigo_unico':code,'estado':status,
                       'fila_oferta':o['fila_origen'] if o else '',
                       'registro_mrun_ejemplo':m['registro_ejemplo'] if m else '',
                       'filas_matricula_mrun':int(m['filas_matricula']) if m else 0,
                       'atributos_diferentes':counts['DIFERENTE'],
                       'atributos_sin_dato':sum(v for k,v in counts.items() if k.startswith('SIN_DATO'))})

    current_comparison=[]
    cy={code:r for (y,code),r in oi.items() if y==current_year}
    for code in sorted(set(ci)|set(cy)):
        o=cy.get(code); c=ci.get(code)
        if not o or not c:
            current_comparison.append({'codigo_unico':code, 'estado':'SOLO_USUARIO' if c else 'SOLO_HISTORICO'})
            continue
        for cf,(of,mapping) in CURRENT_FIELDS.items():
            actual=c[cf]
            if mapping is not None:
                if actual not in mapping:
                    raise ValueError(f'Categoría sin correspondencia gobernada: {cf}={actual}')
                actual=mapping[actual]
            state='IGUAL' if actual==o[of] else 'DIFERENTE'
            current_comparison.append({'codigo_unico':code, 'campo_usuario':cf,'campo_historico':of,
                                       'valor_usuario':c[cf],'valor_usuario_traducido':actual,
                                       'valor_historico':o[of],'estado':state,
                                       'fila_usuario':c['_fila'],'fila_historico':o['fila_origen']})
    by_year=[]
    for y in sorted(set(coverage_years)|{int(r['cat_periodo']) for r in mrun}):
        ar=[r for r in annual if r['anio']==y]
        by_year.append({'anio':y,'oferta_disponible':y in coverage_years,
                        'codigos_oferta':sum(year(r)==y for r in historic),
                        'codigos_mrun':sum(int(r['cat_periodo'])==y for r in mrun),
                        'estados_cruce':dict(Counter(r['estado'] for r in ar))})
    summary={
        'fase':3,'alcance':'Institución 162; reconstrucción de oferta publicada, sin equivalencias UMAS',
        'estado_hito':'RECONSTRUIDO_CON_BRECHA_2007_2009',
        'anios_oferta_disponible':sorted(coverage_years),
        'anios_oferta_pendientes':sorted(set(range(2007,2027))-set(coverage_years)),
        'filas_oferta_anual':len(historic),'codigos_oferta':len(by_code),'familias_oferta':len(families),
        'familias_multiversion':sum(len(f['versiones_informadas'])>1 for f in families),
        'familias_con_varias_versiones_en_primer_anio':sum(len(f['versiones_en_primer_anio'])>1 for f in families),
        'familias_sin_v1_en_primer_anio':[f['familia'] for f in families if '1' not in f['versiones_en_primer_anio']],
        'grupos_familia_anio_coexistentes':len(coexist),'codigos_union_oferta_mrun':len(coverage),
        'presencia_codigos':dict(Counter(r['presencia'] for r in coverage)),
        'codigos_solo_oferta_por_nivel':dict(Counter(r['niveles_globales_oferta'] for r in coverage if r['presencia']=='SOLO_OFERTA')),
        'cruce_anual':dict(Counter(r['estado'] for r in annual)),
        'comparacion_atributos_mrun':dict(comparable),
        'comparacion_2026':dict(Counter(r['estado'] for r in current_comparison)),
        'codigos_oferta_2026':len(ci),'codigos_historicos_sin_oferta_2026':len(set(by_code)-set(ci)),
        'cambios_entre_observaciones_mismo_codigo':len(changes),
        'cambios_por_campo':dict(sorted(Counter(f for r in changes for f in r['cambios']).items())),
        'relacion_inicio_declarado_vs_primera_oferta_por_fila':dict(Counter(r['relacion'] for r in start_relations)),
        'filas_sin_modalidad_oferta':sum(not r['MODALIDAD'] for r in historic),
        'cobertura_por_anio':by_year,
        'limites':[
            'Oferta 2007–2009 no reconstruida; MRUN no sustituye oferta completa.',
            'Primera/última observación y ANO_INICIO no acreditan vigencia de malla UMAS.',
            'VIGENCIA sin estudiantes nuevos no significa cerrado.',
            'Historia publicada armonizada; no garantiza el literal de cada edición original.',
            'Nivel Global de oferta usuario usa esquema histórico 1/3; no valida código de carga CFT/IP2026.',
            'Sin MRUN significa sin matrícula observada en el extracto, no inexistencia de estudiantes.',
            'No se asignaron planes, identidades, equivalencias ni versiones UMAS.',
        ],
    }
    return {'oferta_codigo_anio.tsv':historic, 'historia_codigos.json':histories,
            'versiones_por_familia.json':families, 'coexistencias.json':coexist,
            'cambios_atributos.json':changes,'inicio_declarado_vs_observado.json':start_relations,
            'cobertura_codigos.tsv':coverage,'cruce_oferta_mrun_anual.tsv':annual,
            'comparacion_atributos_mrun.tsv':comparisons,'comparacion_oferta_2026.tsv':current_comparison,
            'ejemplos_academicos.json':{fam:[profile(r) for r in historic if family(r['CODIGO_UNICO'])==fam] for fam in EXAMPLE_FAMILIES},
            'resumen.json':summary}


def tsv_bytes(rows):
    # Unión de campos permite conservar también ausencias en cotejo 2026.
    fields=list(dict.fromkeys(k for r in rows for k in r))
    stream=io.StringIO(newline='')
    writer=csv.DictWriter(stream,fieldnames=fields,delimiter='\t',lineterminator='\n')
    writer.writeheader(); writer.writerows(rows)
    return stream.getvalue().encode()


def safe_write(path, data):
    if path.exists() and path.read_bytes()!=data:
        raise ValueError('Salida existente distinta: '+str(path))
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():path.write_bytes(data)


def read_tsv(path):
    with path.open(encoding='utf-8-sig',newline='') as stream:
        return list(csv.DictReader(stream,delimiter='\t'))


def run(config_path, output):
    from openpyxl import load_workbook
    config=json.loads(config_path.read_text())
    paths={k:Path(v['path']) for k,v in config['sources'].items()}
    for k,p in paths.items():
        if digest(p)!=config['sources'][k]['sha256']:
            raise ValueError('Fuente alterada: '+k)
    phase2=json.loads(paths['manifest_fase2'].read_text())
    if phase2['outputs']['historia_codigo_anio.tsv']!=digest(paths['mrun_anual']):
        raise ValueError('Resumen MRUN no respaldado por manifiesto H2')
    verification=json.loads(paths['verificacion_extraccion'].read_text())
    # Auditoría independiente completa: se valida aquí su contrato común.
    if verification['status']!='PASS' or not all(verification['checks'].values()):
        raise ValueError('Extracción histórica no verificada')
    for k,vk in [('historico_xlsx','original'),('historico_tsv','tsv')]:
        if verification['hashes_after'][vk]!=config['sources'][k]['sha256']:
            raise ValueError('Auditoría de otra fuente: '+k)
    historic=read_tsv(paths['historico_tsv']); mrun=read_tsv(paths['mrun_anual'])
    wb=load_workbook(paths['oferta_2026'],read_only=True,data_only=True)
    try:
        it=wb['Hoja1'].iter_rows(values_only=True);headers=next(it)
        if set(headers)!=set(CURRENT_FIELDS):
            raise ValueError('Columnas no previstas en oferta2026')
        current=[{**{h:text(v) for h,v in zip(headers,values)},'_fila':i}
                 for i,values in enumerate(it,2) if any(v is not None for v in values)]
    finally:wb.close()
    if config['current_year']!=2026 or config['institution']!='162':
        raise ValueError('Este contrato de fase 3 corresponde a institución 162 y oferta actual 2026')
    outputs=build(historic,mrun,current,config['coverage_years'],
                  institution=config['institution'],current_year=config['current_year'])
    data={name:(tsv_bytes(value) if name.endswith('.tsv') else json_bytes(value)) for name,value in outputs.items()}
    data['fuentes.json']=config_path.read_bytes()
    checks={'claves_anuales_unicas':True,'componentes_oferta_coherentes':True,
            'extraccion_original_verificada':True,
            'union_codigos_preservada':len(outputs['cobertura_codigos.tsv'])==len({r['CODIGO_UNICO'] for r in historic}|{r['codigo_unico'] for r in mrun}),
            'filas_matricula_preservadas':sum(int(r['filas_matricula']) for r in mrun)==sum(r['filas_matricula_mrun'] for r in outputs['cruce_oferta_mrun_anual.tsv']),
            'fuentes_sin_modificar':all(digest(p)==config['sources'][k]['sha256'] for k,p in paths.items())}
    if not all(checks.values()):raise ValueError('Control de preservación fallido')
    data['manifest.json']=json_bytes({'schema_version':1,'fase':3,'script_sha256':digest(__file__),
                                    'config_sha256':digest(config_path),'checks':checks,
                                    'outputs':{n:hashlib.sha256(b).hexdigest() for n,b in sorted(data.items())},
                                    'privacidad':'Sin identidades MRUN/RUT/CODCLI; dimensiones académicas, conteos y localizadores.'})
    # Verificar todas las colisiones antes de escribir cualquier archivo.
    for name,b in data.items():
        p=output/name
        if p.exists() and p.read_bytes()!=b:raise ValueError('Salida existente distinta: '+str(p))
    for name,b in data.items():safe_write(output/name,b)
    return outputs['resumen.json']


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--fuentes',type=Path,required=True)
    parser.add_argument('--salida',type=Path,required=True)
    args=parser.parse_args()
    print(json.dumps(run(args.fuentes,args.salida),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
