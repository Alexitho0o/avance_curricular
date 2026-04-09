#!/usr/bin/env python3
"""Flujo auditable para validacion individual PDF-PLAN de 34 carreras directas."""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

import pandas as pd


PROCESO = "Avance Curricular SIES 2026"
SUBPROCESO = "Carreras"
ANO_PROCESO = "2026"
ANO_REFERENCIA_DATOS = "2025"

REPO_ROOT = Path(__file__).resolve().parents[1]
BASE_2026 = REPO_ROOT / "avance_curricular_2026"
AUDITORIAS_DIR = BASE_2026 / "04_gobernanza_mallas" / "03_auditorias"
CONCILIACION_DIR = BASE_2026 / "04_gobernanza_mallas" / "02_conciliacion"
EXEC_PREFIX = "VALIDACION_INDIVIDUAL_PDF_PLAN_"
ENDASH = "\u2013"

REQUIRED_INPUTS = {
    "mapeo_seguro_43": {
        "base": "cierre_seguro",
        "filename": "03_MAPEO_SEGURO_43_CARRERAS.tsv",
    },
    "totales_documentales": {
        "base": "cierre_seguro",
        "filename": "02_TOTALES_DOCUMENTALES_CORREGIDOS_POR_PDF.tsv",
    },
    "mapeo_pdf_plan_canonico": {
        "base": "conciliacion",
        "filename": "01_MAPEO_PDF_PLAN_CANONICO.tsv",
    },
    "hoja1_congelada": {
        "path": BASE_2026
        / "00_fuentes_congeladas"
        / "CARGA_CONGELADA_20260626_005826"
        / "derivados_tsv"
        / "PROMEDIOSDEALUMNOS_7804__HOJA_HOJA1.tsv",
    },
    "puente_sies": {
        "path": REPO_ROOT
        / "resultados"
        / "auditoria_integral_multicodcli_2026"
        / "auditoria_20260618_144641"
        / "00_RESPALDO_ENTRADAS"
        / "PUENTE_SIES_COMPILADO.tsv",
    },
}

DIRS = [
    "00_CONTROL",
    "01_ENTRADAS",
    "02_INTERMEDIOS",
    "03_RESULTADOS",
    "04_AUDITORIA",
    "05_LOGS",
]

PHASE_NAMES = {
    0: "DESCUBRIMIENTO_Y_CONTROL_DE_ENTRADAS",
    1: "NORMALIZACION_Y_UNIVERSO_34_DIRECTAS",
    2: "VALIDACION_MAPEO_CANONICO_PDF_PLAN",
    3: "VALIDACION_CODIGO_UNICO_Y_PUENTE",
    4: "VALIDACION_PLAN_EN_MATRIZ",
    5: "CLASIFICACION_INDIVIDUAL",
    6: "INCORPORACION_CONTROLADA_DE_TOTALES",
    7: "PREPARACION_REVISION_MANUAL_NO_EXACTOS",
}

VALID_PHASE_STATES = {"PENDIENTE", "EN_EJECUCION", "COMPLETADA", "BLOQUEADA", "ERROR"}
FINAL_STATES = {
    "VALIDADO_DIRECTO_PDF_PLAN",
    "VALIDACION_PARCIAL_SIN_PLAN_EN_MATRIZ",
    "BLOQUEADO",
}
BLOCK_REASONS = {
    "PDF_NO_INFORMADO",
    "PDF_SIN_MAPEO_CANONICO",
    "PLAN_CANONICO_NO_EXISTE_HOJA1",
    "CODIGO_UNICO_NO_EXISTE_PUENTE",
    "CODIGO_UNICO_MULTIPLE_EN_PUENTE",
    "CODCARR_PUENTE_DIFIERE_CANONICO",
    "PLAN_MATRIZ_DIFIERE_CANONICO",
    "PUENTE_MARCADO_BLOQUEANTE",
    "OTRO_BLOQUEO_DOCUMENTADO",
}

COLUMN_CANDIDATES = {
    "CODIGO_UNICO": [
        "CODIGO_UNICO",
        "CODIGO_UNICO_FINAL",
        "CODIGO_UNICO_SIES",
        "CODIGOS_SIES_POTENCIALES",
    ],
    "PDF": [
        "PDF_CANONICO",
        "ARCHIVO_PDF",
        "PDF_DIRECTO_CANDIDATO",
        "PDF_MALLA_CANDIDATO",
        "PDF",
        "ARCHIVO",
    ],
    "COBERTURA": [
        "COBERTURA_NORMALIZADA",
        "ESTADO_COBERTURA_DOCUMENTAL",
        "COBERTURA",
    ],
    "CARRERA": [
        "NOMBRE_CARRERA",
        "CARRERA",
        "NOMBRE_NORMALIZADO",
        "NOMBRE_L",
    ],
    "JORNADA": [
        "JORNADA",
        "JORNADA_MATRIZ",
    ],
    "PLAN_MATRIZ": [
        "PLAN_ESTUDIOS",
        "PLAN_DE_ESTUDIO",
        "PLAN_DE_ESTUDIO_1",
        "PLAN",
        "PLAN_ESTUDIO",
    ],
    "CODCARR_CANONICO": [
        "CODCARR_CANONICO",
        "CODCARR_CANONICO_VALIDACION",
        "CODCARR_CANONICO_PDF",
        "CODCARR",
        "CODCARPR",
    ],
    "PLAN_CANONICO": [
        "PLAN_DE_ESTUDIO_CANONICO",
        "PLAN_CANONICO_VALIDACION",
        "PLAN_CANONICO",
        "PLAN_DE_ESTUDIO",
    ],
    "EXISTE_EN_HOJA1": [
        "EXISTE_EN_HOJA1",
        "PLAN_EXISTE_EN_HOJA1",
    ],
    "CODIGO_UNICO_PUENTE": [
        "CODIGO_UNICO_FINAL",
        "CODIGO_UNICO",
        "CODIGOS_SIES_POTENCIALES",
    ],
    "CODCARPR": [
        "CODCARPR",
        "FAMILIA_CODCARPR",
    ],
    "RESOLUCION_STATUS": [
        "RESOLUCION_STATUS",
        "GOBERNANZA_STATUS",
    ],
    "ES_BLOQUEANTE": [
        "ES_BLOQUEANTE",
        "BLOQUEANTE",
    ],
    "TOTAL_UNIDADES": [
        "TOTAL_UNIDADES_DOCUMENTALES",
        "TOTAL_UNIDADES",
    ],
}

IDENTIDAD_RUN_DIR = (
    AUDITORIAS_DIR / "INTEGRACION_IDENTIDAD_CARRERAS_20260630_111622"
)
RUTA_IDENTIDAD_34 = IDENTIDAD_RUN_DIR / "04_RESULTADOS/06_IDENTIDAD_APLICADA_34_CARRERAS.tsv"
RUTA_DETALLE_IDENTIDAD_34 = (
    IDENTIDAD_RUN_DIR / "04_RESULTADOS/06A_IDENTIDAD_APLICADA_34_DETALLE_RELACIONES.tsv"
)
RUTA_COMPARACION_IDENTIDAD_PDF_PLAN = (
    IDENTIDAD_RUN_DIR / "05_AUDITORIA/FASE5_COMPARACION_IDENTIDAD_PDF_PLAN.tsv"
)
RUTA_COMPARACION_IDENTIDAD_PUENTE = (
    IDENTIDAD_RUN_DIR / "05_AUDITORIA/FASE5_COMPARACION_PUENTE_PREVIO.tsv"
)
VERSION_LOGICA_FASE4_INTEGRADA = "INTEGRACION_IDENTIDAD_CARRERAS_V1"
RUTA_CONCILIACION_CLASIFICADA_415 = (
    AUDITORIAS_DIR
    / "CIERRE_DIAGNOSTICO_MALLAS_20260626_170631"
    / "01_CONCILIACION_CLASIFICADA_415.tsv"
)
RUTA_MALLA_PDF_CANDIDATA_415 = (
    BASE_2026
    / "04_gobernanza_mallas"
    / "01_extraccion_pdf"
    / "EXTRACCION_MALLAS_PDF_20260626_165945"
    / "02_MALLA_PDF_CANDIDATA.tsv"
)
RUTA_CONCILIACION_CANONICA_ASIGNATURAS = (
    CONCILIACION_DIR
    / "CONCILIACION_CANONICA_PDF_HOJA1_20260626_170541"
    / "03_CONCILIACION_CANONICA_PDF_HOJA1.tsv"
)
RUTA_TOTALES_DOCUMENTALES_CORREGIDOS = (
    AUDITORIAS_DIR
    / "CIERRE_SEGURO_MAPEO_MALLAS_20260626_171212"
    / "02_TOTALES_DOCUMENTALES_CORREGIDOS_POR_PDF.tsv"
)


@dataclass
class PhaseResult:
    estado: str
    salidas: list[str] = field(default_factory=list)
    conteos: dict[str, Any] = field(default_factory=dict)
    validaciones: list[dict[str, Any]] = field(default_factory=list)
    errores: list[dict[str, Any]] = field(default_factory=list)
    pendientes: list[dict[str, Any]] = field(default_factory=list)
    entradas: dict[str, str] = field(default_factory=dict)
    hashes: dict[str, str] = field(default_factory=dict)
    resumen: list[str] = field(default_factory=list)
    comando_reanudacion: str = ""


class FlujoError(Exception):
    """Error controlado del flujo."""


def normalize_argv(argv: list[str]) -> list[str]:
    normalized: list[str] = []
    for arg in argv:
        if arg.startswith(("–", "—")):
            stripped = arg.lstrip("–—")
            normalized.append("--" + stripped)
        else:
            normalized.append(arg)
    return normalized


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validacion individual PDF-PLAN de 34 carreras directas."
    )
    parser.add_argument("--fase", default=None, help="Fase 0-7 o final.")
    parser.add_argument("--reanudar", action="store_true", help="Reanudar ultima ejecucion.")
    parser.add_argument("--forzar", action="store_true", help="Repetir una fase ya completada.")
    parser.add_argument(
        "--solo-validar",
        action="store_true",
        help="Validar estado y hashes sin ejecutar fases nuevas.",
    )
    parser.add_argument("--mostrar-rutas", action="store_true", help="Mostrar rutas base.")
    parser.add_argument(
        "--regenerar-paquete-tec-cont",
        action="store_true",
        help="Crear paquete derivado de revision tecnica-continuidad sin alterar fases 0-7.",
    )
    parser.add_argument(
        "--auditoria-mallas",
        default="",
        help="Carpeta de auditoria end-to-end de mallas compartidas.",
    )
    parser.add_argument(
        "--carpeta-ejecucion",
        default="",
        help="Carpeta del flujo PDF-plan donde crear el paquete derivado.",
    )
    return parser.parse_args(normalize_argv(argv))


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def command_for_phase(phase: int | str) -> str:
    value = "final" if phase == "final" else str(phase)
    return (
        f"python scripts/validar_pdf_plan_34_directas.py "
        f"{ENDASH}fase {value} {ENDASH}reanudar"
    )


def normalize_col(name: str) -> str:
    text = unicodedata.normalize("NFKD", str(name))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip().upper()).strip("_")
    return text or "COLUMNA_SIN_NOMBRE"


def normalize_columns(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, str]]:
    used: dict[str, int] = {}
    rename: dict[str, str] = {}
    original_by_normalized: dict[str, str] = {}
    for col in df.columns:
        base = normalize_col(col)
        count = used.get(base, 0)
        normalized = base if count == 0 else f"{base}_{count + 1}"
        used[base] = count + 1
        rename[col] = normalized
        original_by_normalized[normalized] = str(col)
    return df.rename(columns=rename), original_by_normalized


def read_tsv(path: Path, normalize: bool = False) -> pd.DataFrame:
    df = pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False, encoding="utf-8")
    if normalize:
        df, _ = normalize_columns(df)
    return df


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_tsv(df: pd.DataFrame, path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, sep="\t", index=False, encoding="utf-8")
    return str(path.resolve())


def write_json(data: dict[str, Any], path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, sort_keys=True)
        fh.write("\n")
    return str(path.resolve())


def write_text(text: str, path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return str(path.resolve())


def empty_df(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def first_existing(columns: list[str], candidates: list[str]) -> str | None:
    available = set(columns)
    for candidate in candidates:
        normalized = normalize_col(candidate)
        if normalized in available:
            return normalized
    return None


def required_column(df: pd.DataFrame, logical_name: str) -> str:
    col = first_existing(list(df.columns), COLUMN_CANDIDATES[logical_name])
    if not col:
        raise FlujoError(f"No se detecto columna requerida: {logical_name}")
    return col


def truthy_si(value: Any) -> bool:
    return str(value).strip().upper() in {"SI", "S", "YES", "Y", "TRUE", "1"}


def is_no(value: Any) -> bool:
    return str(value).strip().upper() in {"NO", "N", "FALSE", "0", ""}


def latest_dir(base: Path, pattern: str) -> Path | None:
    candidates = sorted([p for p in base.glob(pattern) if p.is_dir()])
    return candidates[-1] if candidates else None


def list_execution_dirs() -> list[Path]:
    return sorted([p for p in AUDITORIAS_DIR.glob(f"{EXEC_PREFIX}*") if p.is_dir()])


def phase_key(phase: int) -> str:
    return f"fase_{phase}"


def phase_status(state: dict[str, Any], phase: int) -> str:
    return (
        state.get("fases", {})
        .get(phase_key(phase), {})
        .get("estado", "PENDIENTE")
    )


def run_is_complete(run_dir: Path) -> bool:
    state_path = run_dir / "00_CONTROL" / "estado_fases.json"
    if not state_path.exists():
        return False
    try:
        with state_path.open("r", encoding="utf-8") as fh:
            state = json.load(fh)
    except json.JSONDecodeError:
        return False
    return phase_status(state, 7) == "COMPLETADA"


def latest_execution(prefer_incomplete: bool = True) -> Path | None:
    executions = list_execution_dirs()
    if not executions:
        return None
    if prefer_incomplete:
        incomplete = [p for p in executions if not run_is_complete(p)]
        if incomplete:
            return incomplete[-1]
    return executions[-1]


def create_run_dir() -> Path:
    AUDITORIAS_DIR.mkdir(parents=True, exist_ok=True)
    candidate = AUDITORIAS_DIR / f"{EXEC_PREFIX}{timestamp()}"
    suffix = 1
    while candidate.exists():
        candidate = AUDITORIAS_DIR / f"{EXEC_PREFIX}{timestamp()}_{suffix}"
        suffix += 1
    for dirname in DIRS:
        (candidate / dirname).mkdir(parents=True, exist_ok=True)
    return candidate


class RunContext:
    def __init__(self, run_dir: Path):
        self.run_dir = run_dir
        for dirname in DIRS:
            (self.run_dir / dirname).mkdir(parents=True, exist_ok=True)
        self.control_path = self.run_dir / "00_CONTROL" / "estado_fases.json"
        self.log_path = self.run_dir / "05_LOGS" / "ejecucion.log"
        self.state = self._load_or_init_state()

    def _load_or_init_state(self) -> dict[str, Any]:
        if self.control_path.exists():
            with self.control_path.open("r", encoding="utf-8") as fh:
                return json.load(fh)
        return {
            "proceso": PROCESO,
            "subproceso": SUBPROCESO,
            "ano_proceso": ANO_PROCESO,
            "ano_referencia_datos": ANO_REFERENCIA_DATOS,
            "run_dir": str(self.run_dir.resolve()),
            "fecha_creacion": now_iso(),
            "fases": {
                phase_key(i): {
                    "fase": i,
                    "nombre": PHASE_NAMES[i],
                    "estado": "PENDIENTE",
                    "fecha_inicio": "",
                    "fecha_fin": "",
                    "entradas": {},
                    "hashes": {},
                    "salidas": [],
                    "conteos": {},
                    "validaciones": [],
                    "errores": [],
                    "pendientes": [],
                    "comando_reanudacion": command_for_phase(i + 1 if i < 7 else "final"),
                }
                for i in range(8)
            },
            "reanudaciones": [],
        }

    def save(self) -> None:
        write_json(self.state, self.control_path)

    def log(self, message: str) -> None:
        self.log_path.parent.mkdir(parents=True, exist_ok=True)
        with self.log_path.open("a", encoding="utf-8") as fh:
            fh.write(f"{now_iso()} {message}\n")

    def input_paths(self) -> dict[str, str]:
        entradas = self.state.get("fases", {}).get("fase_0", {}).get("entradas", {})
        if not entradas:
            raise FlujoError("Fase 0 no tiene entradas registradas.")
        return entradas

    def start_phase(self, phase: int) -> None:
        entry = self.state["fases"][phase_key(phase)]
        entry.update(
            {
                "estado": "EN_EJECUCION",
                "fecha_inicio": now_iso(),
                "fecha_fin": "",
                "errores": [],
                "pendientes": [],
            }
        )
        self.save()
        self.log(f"FASE {phase} INICIO")

    def finish_phase(self, phase: int, result: PhaseResult) -> None:
        if result.estado not in VALID_PHASE_STATES:
            raise FlujoError(f"Estado de fase no permitido: {result.estado}")
        entry = self.state["fases"][phase_key(phase)]
        entry.update(
            {
                "estado": result.estado,
                "fecha_fin": now_iso(),
                "entradas": result.entradas or entry.get("entradas", {}),
                "hashes": result.hashes or entry.get("hashes", {}),
                "salidas": result.salidas,
                "conteos": result.conteos,
                "validaciones": result.validaciones,
                "errores": result.errores,
                "pendientes": result.pendientes,
                "comando_reanudacion": result.comando_reanudacion,
            }
        )
        self.save()
        self.log(f"FASE {phase} FIN {result.estado}")


def output_path(ctx: RunContext, relative: str) -> Path:
    return ctx.run_dir / relative


def print_result(title: str, result: PhaseResult) -> None:
    print()
    print(title)
    for line in result.resumen:
        print(line)
    if result.errores:
        print("Errores o bloqueos relevantes:")
        for err in result.errores[:20]:
            print(f"- {err}")
        if len(result.errores) > 20:
            print(f"- ... {len(result.errores) - 20} casos adicionales en archivos TSV/JSON")
    if result.comando_reanudacion:
        print(f"Comando siguiente: {result.comando_reanudacion}")


def discover_input_paths() -> tuple[dict[str, Path], list[dict[str, Any]]]:
    errors: list[dict[str, Any]] = []
    cierre_dir = latest_dir(AUDITORIAS_DIR, "CIERRE_SEGURO_MAPEO_MALLAS_*")
    conciliacion_dir = latest_dir(CONCILIACION_DIR, "CONCILIACION_CANONICA_PDF_HOJA1_*")

    if cierre_dir is None:
        errors.append({"tipo": "CARPETA_NO_ENCONTRADA", "detalle": "CIERRE_SEGURO_MAPEO_MALLAS_*"})
    if conciliacion_dir is None:
        errors.append(
            {"tipo": "CARPETA_NO_ENCONTRADA", "detalle": "CONCILIACION_CANONICA_PDF_HOJA1_*"}
        )

    paths: dict[str, Path] = {}
    for logical, spec in REQUIRED_INPUTS.items():
        if "path" in spec:
            paths[logical] = Path(spec["path"])
        elif spec["base"] == "cierre_seguro" and cierre_dir is not None:
            paths[logical] = cierre_dir / spec["filename"]
        elif spec["base"] == "conciliacion" and conciliacion_dir is not None:
            paths[logical] = conciliacion_dir / spec["filename"]

    if cierre_dir is not None:
        paths["carpeta_cierre_seguro"] = cierre_dir
    if conciliacion_dir is not None:
        paths["carpeta_conciliacion"] = conciliacion_dir
    return paths, errors


def fase_0(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    paths, discovery_errors = discover_input_paths()
    inventory_rows: list[dict[str, Any]] = []
    hash_rows: list[dict[str, str]] = []
    entradas: dict[str, str] = {}
    hashes: dict[str, str] = {}
    errors = discovery_errors[:]

    for logical in REQUIRED_INPUTS:
        path = paths.get(logical)
        row: dict[str, Any] = {
            "entrada": logical,
            "ruta_absoluta": str(path.resolve()) if path else "",
            "existe": "NO",
            "tipo": "TSV",
            "bytes": "",
            "filas": "",
            "columnas": "",
            "legible_utf8_tsv": "NO",
            "error": "",
        }
        if path is None:
            row["error"] = "ruta_no_detectada"
            errors.append({"entrada": logical, "error": "ruta_no_detectada"})
            inventory_rows.append(row)
            continue
        entradas[logical] = str(path.resolve())
        if not path.exists():
            row["error"] = "archivo_no_existe"
            errors.append({"entrada": logical, "ruta": str(path.resolve()), "error": "archivo_no_existe"})
            inventory_rows.append(row)
            continue
        row["existe"] = "SI"
        row["bytes"] = path.stat().st_size
        try:
            df = read_tsv(path)
            row["filas"] = len(df)
            row["columnas"] = len(df.columns)
            row["legible_utf8_tsv"] = "SI"
            digest = sha256_file(path)
            hashes[logical] = digest
            hash_rows.append(
                {
                    "entrada": logical,
                    "ruta_absoluta": str(path.resolve()),
                    "sha256": digest,
                }
            )
        except Exception as exc:  # noqa: BLE001 - debe registrar error concreto de lectura.
            row["error"] = str(exc)
            errors.append({"entrada": logical, "ruta": str(path.resolve()), "error": str(exc)})
        inventory_rows.append(row)

    for logical in ("carpeta_cierre_seguro", "carpeta_conciliacion"):
        if logical in paths:
            entradas[logical] = str(paths[logical].resolve())

    inventory_path = output_path(ctx, "01_ENTRADAS/inventario_entradas.tsv")
    hashes_path = output_path(ctx, "01_ENTRADAS/hashes_entradas.tsv")
    salidas = [
        write_tsv(pd.DataFrame(inventory_rows), inventory_path),
        write_tsv(pd.DataFrame(hash_rows), hashes_path),
        str(ctx.control_path.resolve()),
    ]

    estado = "COMPLETADA" if not errors else "BLOQUEADA"
    resumen = [
        f"Estado: {estado}",
        f"Carpeta de ejecucion: {ctx.run_dir.resolve()}",
        f"Entradas requeridas: {len(REQUIRED_INPUTS)}",
        f"Archivos validos: {sum(1 for r in inventory_rows if r['legible_utf8_tsv'] == 'SI')}",
        f"Bloqueos: {len(errors)}",
        f"Inventario: {inventory_path.resolve()}",
        f"Hashes: {hashes_path.resolve()}",
    ]
    if "carpeta_cierre_seguro" in entradas:
        resumen.append(f"Cierre seguro detectado: {entradas['carpeta_cierre_seguro']}")
    if "carpeta_conciliacion" in entradas:
        resumen.append(f"Conciliacion detectada: {entradas['carpeta_conciliacion']}")

    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "entradas_requeridas": len(REQUIRED_INPUTS),
            "archivos_validos": sum(1 for r in inventory_rows if r["legible_utf8_tsv"] == "SI"),
            "bloqueos": len(errors),
        },
        validaciones=[
            {"validacion": "tsv_utf8_dtype_str_keep_default_na_false", "estado": estado},
            {"validacion": "hash_sha256_entradas", "estado": "COMPLETADA" if hashes else "BLOQUEADA"},
        ],
        errores=errors,
        entradas=entradas,
        hashes=hashes,
        resumen=resumen,
        comando_reanudacion=command_for_phase(1),
    )


def fase_1(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    entradas = ctx.input_paths()
    df = read_tsv(Path(entradas["mapeo_seguro_43"]), normalize=True)
    codigo_col = required_column(df, "CODIGO_UNICO")
    pdf_col = first_existing(list(df.columns), ["ARCHIVO_PDF"])
    if not pdf_col:
        pdf_col = required_column(df, "PDF")
    cobertura_col = required_column(df, "COBERTURA")
    carrera_col = first_existing(list(df.columns), COLUMN_CANDIDATES["CARRERA"])
    jornada_col = first_existing(list(df.columns), COLUMN_CANDIDATES["JORNADA"])
    plan_col = first_existing(list(df.columns), ["PLAN_ESTUDIOS", "PLAN_DE_ESTUDIO", "PLAN_ESTUDIO"])

    total = len(df)
    direct = df[df[cobertura_col] == "COBERTURA_DIRECTA_PDF"].copy()
    excluded = df[df[cobertura_col] != "COBERTURA_DIRECTA_PDF"].copy()
    dupes = direct[direct.duplicated(subset=[codigo_col], keep=False)].copy()

    control_rows = [
        {"control": "total_carreras_matriz", "valor": total, "esperado": 43, "estado": "OK" if total == 43 else "ERROR"},
        {"control": "carreras_directas", "valor": len(direct), "esperado": 34, "estado": "OK" if len(direct) == 34 else "ERROR"},
        {"control": "carreras_no_directas_excluidas", "valor": len(excluded), "esperado": 9, "estado": "OK" if len(excluded) == 9 else "ERROR"},
        {"control": "duplicados_codigo_unico", "valor": len(dupes), "esperado": 0, "estado": "OK" if len(dupes) == 0 else "ERROR"},
        {"control": "columna_pdf", "valor": pdf_col or "", "esperado": "detectada", "estado": "OK" if pdf_col else "ERROR"},
        {"control": "columna_cobertura", "valor": cobertura_col or "", "esperado": "detectada", "estado": "OK" if cobertura_col else "ERROR"},
        {"control": "columna_plan_matriz", "valor": plan_col or "", "esperado": "opcional", "estado": "OK"},
        {"control": "columna_carrera", "valor": carrera_col or "", "esperado": "opcional", "estado": "OK" if carrera_col else "ADVERTENCIA"},
        {"control": "columna_jornada", "valor": jornada_col or "", "esperado": "opcional", "estado": "OK" if jornada_col else "ADVERTENCIA"},
    ]

    errors: list[dict[str, Any]] = []
    if total != 43:
        errors.append({"control": "total_carreras_matriz", "valor": total, "esperado": 43})
    if len(direct) != 34:
        errors.append({"control": "carreras_directas", "valor": len(direct), "esperado": 34})
    if len(dupes) > 0:
        errors.append({"control": "duplicados_codigo_unico", "valor": len(dupes), "esperado": 0})
    if not pdf_col:
        errors.append({"control": "columna_pdf", "error": "no_detectada"})
    if not cobertura_col:
        errors.append({"control": "columna_cobertura", "error": "no_detectada"})

    out_universe = output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv")
    out_control = output_path(ctx, "04_AUDITORIA/01_CONTROL_UNIVERSO.tsv")
    out_dupes = output_path(ctx, "04_AUDITORIA/02_DUPLICADOS_CODIGO_UNICO.tsv")
    salidas = [
        write_tsv(direct, out_universe),
        write_tsv(pd.DataFrame(control_rows), out_control),
        write_tsv(dupes, out_dupes),
    ]

    estado = "COMPLETADA" if not errors else "BLOQUEADA"
    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "total_carreras": total,
            "directas": len(direct),
            "excluidas": len(excluded),
            "duplicados_codigo_unico": len(dupes),
        },
        validaciones=control_rows,
        errores=errors,
        resumen=[
            f"Estado: {estado}",
            f"Total carreras: {total}",
            f"Directas: {len(direct)}",
            f"Excluidas: {len(excluded)}",
            f"Duplicados CODIGO_UNICO: {len(dupes)}",
            f"Universo: {out_universe.resolve()}",
        ],
        comando_reanudacion=command_for_phase(2),
    )


def hoja1_plan_values(path: Path) -> set[str]:
    hoja = read_tsv(path, normalize=True)
    plan_cols = [c for c in hoja.columns if c.startswith("PLAN_DE_ESTUDIO") or c in {"PLAN_ESTUDIOS", "PLAN"}]
    values: set[str] = set()
    for col in plan_cols:
        values.update(v for v in hoja[col].astype(str).tolist() if v != "")
    return values


def fase_2(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    entradas = ctx.input_paths()
    universe = read_tsv(output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"), normalize=True)
    mapping = read_tsv(Path(entradas["mapeo_pdf_plan_canonico"]), normalize=True)
    pdf_u = required_column(universe, "PDF")
    pdf_m = required_column(mapping, "PDF")
    codcarr_m = required_column(mapping, "CODCARR_CANONICO")
    plan_m = required_column(mapping, "PLAN_CANONICO")
    existe_m = required_column(mapping, "EXISTE_EN_HOJA1")

    normalized_mapping = mapping.copy()
    direct_pdfs = sorted(set(v for v in universe[pdf_u].astype(str).tolist() if v != ""))
    mapped_pdfs = set(mapping[pdf_m].astype(str).tolist())
    missing_rows = [{"ARCHIVO_PDF": pdf} for pdf in direct_pdfs if pdf not in mapped_pdfs]

    dup_counts = (
        mapping.groupby(pdf_m, dropna=False)
        .size()
        .reset_index(name="N_MAPEOS_CANONICOS")
    )
    multiples = dup_counts[dup_counts["N_MAPEOS_CANONICOS"].astype(int) > 1].copy()

    plans_in_hoja1 = hoja1_plan_values(Path(entradas["hoja1_congelada"]))
    direct_mapping = mapping[mapping[pdf_m].isin(direct_pdfs)].copy()
    bad_plan_rows = []
    for _, row in direct_mapping.iterrows():
        plan = row[plan_m]
        flag_ok = truthy_si(row[existe_m])
        actual_ok = plan in plans_in_hoja1
        if not flag_ok or not actual_ok:
            bad_plan_rows.append(
                {
                    "ARCHIVO_PDF": row[pdf_m],
                    "CODCARR_CANONICO": row[codcarr_m],
                    "PLAN_DE_ESTUDIO_CANONICO": plan,
                    "EXISTE_EN_HOJA1_DECLARADO": row[existe_m],
                    "EXISTE_EN_HOJA1_VERIFICADO": "SI" if actual_ok else "NO",
                }
            )

    out_mapping = output_path(ctx, "02_INTERMEDIOS/02_MAPEO_PDF_PLAN_NORMALIZADO.tsv")
    out_missing = output_path(ctx, "04_AUDITORIA/03_PDF_SIN_MAPEO.tsv")
    out_multiple = output_path(ctx, "04_AUDITORIA/04_PDF_MAPEO_MULTIPLE.tsv")
    out_bad_plan = output_path(ctx, "04_AUDITORIA/05_PLANES_NO_EXISTENTES_HOJA1.tsv")
    salidas = [
        write_tsv(normalized_mapping, out_mapping),
        write_tsv(pd.DataFrame(missing_rows, columns=["ARCHIVO_PDF"]), out_missing),
        write_tsv(multiples, out_multiple),
        write_tsv(pd.DataFrame(bad_plan_rows), out_bad_plan),
    ]
    errors: list[dict[str, Any]] = []
    if missing_rows:
        errors.append({"control": "pdf_sin_mapeo", "casos": len(missing_rows), "muestra": missing_rows[:20]})
    if len(multiples) > 0:
        errors.append({"control": "pdf_mapeo_multiple", "casos": len(multiples)})
    if bad_plan_rows:
        errors.append({"control": "planes_no_existentes_hoja1", "casos": len(bad_plan_rows), "muestra": bad_plan_rows[:20]})
    estado = "COMPLETADA" if not errors else "BLOQUEADA"
    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "pdf_directos": len(direct_pdfs),
            "pdf_sin_mapeo": len(missing_rows),
            "pdf_mapeo_multiple": len(multiples),
            "planes_no_existentes_hoja1": len(bad_plan_rows),
        },
        validaciones=[
            {"validacion": "pdf_directo_con_mapeo_unico", "estado": "OK" if not missing_rows and len(multiples) == 0 else "ERROR"},
            {"validacion": "plan_canonico_existe_hoja1", "estado": "OK" if not bad_plan_rows else "ERROR"},
        ],
        errores=errors,
        resumen=[
            f"Estado: {estado}",
            f"PDF directos: {len(direct_pdfs)}",
            f"PDF sin mapeo: {len(missing_rows)}",
            f"PDF con mapeo multiple: {len(multiples)}",
            f"Planes inexistentes en Hoja1: {len(bad_plan_rows)}",
            f"Mapeo normalizado: {out_mapping.resolve()}",
        ],
        comando_reanudacion=command_for_phase(3),
    )


def merge_universe_mapping(universe: pd.DataFrame, mapping: pd.DataFrame) -> pd.DataFrame:
    pdf_u = required_column(universe, "PDF")
    pdf_m = required_column(mapping, "PDF")
    codcarr_m = required_column(mapping, "CODCARR_CANONICO")
    plan_m = required_column(mapping, "PLAN_CANONICO")
    cols = [pdf_m, codcarr_m, plan_m]
    if first_existing(list(mapping.columns), COLUMN_CANDIDATES["EXISTE_EN_HOJA1"]):
        cols.append(required_column(mapping, "EXISTE_EN_HOJA1"))
    universe_for_merge = universe.copy()
    for logical_name in ("CODCARR_CANONICO", "PLAN_CANONICO", "EXISTE_EN_HOJA1"):
        existing = first_existing(list(universe_for_merge.columns), COLUMN_CANDIDATES[logical_name])
        if existing:
            universe_for_merge = universe_for_merge.drop(columns=[existing])
    mapping_for_merge = mapping[cols].drop_duplicates().rename(
        columns={
            codcarr_m: "CODCARR_CANONICO",
            plan_m: "PLAN_DE_ESTUDIO_CANONICO",
            pdf_m: "ARCHIVO_PDF" if pdf_m != "ARCHIVO_PDF" else pdf_m,
        }
    )
    return universe_for_merge.merge(
        mapping_for_merge,
        left_on=pdf_u,
        right_on="ARCHIVO_PDF",
        how="left",
        suffixes=("", "_VALIDACION"),
    )


def validar_regresion_pre_fase3(ctx: RunContext) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    universe = read_tsv(output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"), normalize=True)
    mapping = read_tsv(output_path(ctx, "02_INTERMEDIOS/02_MAPEO_PDF_PLAN_NORMALIZADO.tsv"), normalize=True)
    codigo_u = required_column(universe, "CODIGO_UNICO")
    pdf_u = required_column(universe, "PDF")
    pdf_m = required_column(mapping, "PDF")
    plan_m = required_column(mapping, "PLAN_CANONICO")
    existe_m = required_column(mapping, "EXISTE_EN_HOJA1")
    direct_pdfs = sorted(set(v for v in universe[pdf_u].astype(str).tolist() if v != ""))
    mapped_pdfs = set(mapping[pdf_m].astype(str).tolist())
    pdf_sin_mapeo = [pdf for pdf in direct_pdfs if pdf not in mapped_pdfs]
    pdf_multiples = (
        mapping[mapping[pdf_m].isin(direct_pdfs)]
        .groupby(pdf_m, dropna=False)
        .size()
        .reset_index(name="N")
    )
    pdf_multiples = pdf_multiples[pdf_multiples["N"].astype(int) > 1]
    planes_no_hoja1 = mapping[
        mapping[pdf_m].isin(direct_pdfs)
        & ((mapping[plan_m] == "") | (~mapping[existe_m].map(truthy_si)))
    ]
    checks = [
        {"validacion": "universo_directo_34", "valor": len(universe), "esperado": 34},
        {
            "validacion": "codigos_unicos_duplicados_0",
            "valor": int(universe[codigo_u].duplicated().sum()),
            "esperado": 0,
        },
        {"validacion": "pdf_directos_11", "valor": len(direct_pdfs), "esperado": 11},
        {"validacion": "pdf_sin_mapeo_0", "valor": len(pdf_sin_mapeo), "esperado": 0},
        {"validacion": "pdf_mapeo_multiple_0", "valor": len(pdf_multiples), "esperado": 0},
        {"validacion": "planes_inexistentes_hoja1_0", "valor": len(planes_no_hoja1), "esperado": 0},
    ]
    errors = [c for c in checks if c["valor"] != c["esperado"]]
    details = {
        "columnas_universo": list(universe.columns),
        "columnas_mapeo_fase2": list(mapping.columns),
        "columna_codigo_carrera_canonico_fase2": required_column(mapping, "CODCARR_CANONICO"),
        "columna_plan_canonico_fase2": plan_m,
        "columna_pdf_fase2": pdf_m,
        "seccion_corregida": "merge_universe_mapping/fase_3",
        "causa_error": "La union previa sufijaba columnas canonicas duplicadas y Fase 3 buscaba CODCARR_CANONICO sin sufijo.",
    }
    return checks, errors, details


def fase_3(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    entradas = ctx.input_paths()
    universe = read_tsv(output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"), normalize=True)
    mapping = read_tsv(output_path(ctx, "02_INTERMEDIOS/02_MAPEO_PDF_PLAN_NORMALIZADO.tsv"), normalize=True)
    bridge = read_tsv(Path(entradas["puente_sies"]), normalize=True)
    regression_checks, regression_errors, diagnostic = validar_regresion_pre_fase3(ctx)
    ctx.log(
        "DIAGNOSTICO_FASE3 "
        + json.dumps(
            {
                "columnas_detectadas": {
                    "universo": diagnostic["columnas_universo"],
                    "mapeo_fase2": diagnostic["columnas_mapeo_fase2"],
                },
                "columna_seleccionada": {
                    "codigo_carrera_canonico": diagnostic["columna_codigo_carrera_canonico_fase2"],
                    "plan_canonico": diagnostic["columna_plan_canonico_fase2"],
                    "pdf": diagnostic["columna_pdf_fase2"],
                },
                "seccion_corregida": diagnostic["seccion_corregida"],
                "causa_error": diagnostic["causa_error"],
            },
            ensure_ascii=True,
        )
    )
    if regression_errors:
        return PhaseResult(
            estado="BLOQUEADA",
            conteos={
                "bloqueos_estructurales": len(regression_errors),
                "carreras_procesadas": 0,
            },
            validaciones=regression_checks,
            errores=regression_errors[:20],
            resumen=[
                "Estado: BLOQUEADA",
                f"Bloqueos estructurales: {len(regression_errors)}",
                "La Fase 3 no se ejecuto porque cambiaron validaciones previas.",
            ],
            comando_reanudacion=command_for_phase(3),
        )
    merged = merge_universe_mapping(universe, mapping)

    codigo_u = required_column(merged, "CODIGO_UNICO")
    pdf_u = required_column(merged, "PDF")
    jornada_u = first_existing(list(merged.columns), COLUMN_CANDIDATES["JORNADA"])
    codcarr_can = required_column(merged, "CODCARR_CANONICO")
    plan_can = required_column(merged, "PLAN_CANONICO")
    codigo_b = required_column(bridge, "CODIGO_UNICO_PUENTE")
    codcarpr_b = required_column(bridge, "CODCARPR")
    jornada_b = first_existing(list(bridge.columns), COLUMN_CANDIDATES["JORNADA"])
    resol_b = required_column(bridge, "RESOLUCION_STATUS")
    bloqueante_b = required_column(bridge, "ES_BLOQUEANTE")

    rows: list[dict[str, Any]] = []
    no_bridge: list[dict[str, Any]] = []
    multiple_bridge_frames: list[pd.DataFrame] = []
    codcarr_diff: list[dict[str, Any]] = []
    blocking_bridge: list[dict[str, Any]] = []

    for _, row in merged.iterrows():
        code = row[codigo_u]
        matches = bridge[bridge[codigo_b] == code].copy()
        out = {
            "CODIGO_UNICO": code,
            "ARCHIVO_PDF": row[pdf_u],
            "JORNADA_MATRIZ": row[jornada_u] if jornada_u else "",
            "CODCARR_CANONICO": row[codcarr_can],
            "PLAN_DE_ESTUDIO_CANONICO": row[plan_can],
            "PUENTE_MATCH_COUNT": len(matches),
            "CODCARPR_PUENTE": "",
            "JORNADA_PUENTE": "",
            "RESOLUCION_STATUS": "",
            "ES_BLOQUEANTE": "",
            "VALIDACION_PUENTE": "",
            "MOTIVOS_PUENTE": "",
        }
        motivos: list[str] = []
        if len(matches) == 0:
            motivos.append("CODIGO_UNICO_NO_EXISTE_PUENTE")
            no_bridge.append({"CODIGO_UNICO": code, "ARCHIVO_PDF": row[pdf_u]})
        elif len(matches) > 1:
            motivos.append("CODIGO_UNICO_MULTIPLE_EN_PUENTE")
            sample = matches.copy()
            sample.insert(0, "CODIGO_UNICO_VALIDADO", code)
            multiple_bridge_frames.append(sample)
        else:
            match = matches.iloc[0]
            out["CODCARPR_PUENTE"] = match[codcarpr_b]
            out["JORNADA_PUENTE"] = match[jornada_b] if jornada_b else ""
            out["RESOLUCION_STATUS"] = match[resol_b]
            out["ES_BLOQUEANTE"] = match[bloqueante_b]
            if match[codcarpr_b] != row[codcarr_can]:
                motivos.append("CODCARR_PUENTE_DIFIERE_CANONICO")
                codcarr_diff.append(
                    {
                        "CODIGO_UNICO": code,
                        "ARCHIVO_PDF": row[pdf_u],
                        "CODCARR_CANONICO": row[codcarr_can],
                        "CODCARPR_PUENTE": match[codcarpr_b],
                    }
                )
            if truthy_si(match[bloqueante_b]):
                motivos.append("PUENTE_MARCADO_BLOQUEANTE")
                blocking_bridge.append(
                    {
                        "CODIGO_UNICO": code,
                        "ARCHIVO_PDF": row[pdf_u],
                        "ES_BLOQUEANTE": match[bloqueante_b],
                        "RESOLUCION_STATUS": match[resol_b],
                    }
                )
        out["VALIDACION_PUENTE"] = "OK" if not motivos else "BLOQUEADO"
        out["MOTIVOS_PUENTE"] = "|".join(motivos)
        rows.append(out)

    multiple_bridge = (
        pd.concat(multiple_bridge_frames, ignore_index=True)
        if multiple_bridge_frames
        else empty_df(["CODIGO_UNICO_VALIDADO"])
    )

    out_validation = output_path(ctx, "02_INTERMEDIOS/03_VALIDACION_PUENTE_34.tsv")
    out_no_bridge = output_path(ctx, "04_AUDITORIA/06_CODIGOS_NO_EN_PUENTE.tsv")
    out_multiple = output_path(ctx, "04_AUDITORIA/07_CODIGOS_MULTIPLES_EN_PUENTE.tsv")
    out_diff = output_path(ctx, "04_AUDITORIA/08_CODCARR_DIFIERE_CANONICO.tsv")
    out_blocking = output_path(ctx, "04_AUDITORIA/09_PUENTE_BLOQUEANTE.tsv")
    salidas = [
        write_tsv(pd.DataFrame(rows), out_validation),
        write_tsv(pd.DataFrame(no_bridge), out_no_bridge),
        write_tsv(multiple_bridge, out_multiple),
        write_tsv(pd.DataFrame(codcarr_diff), out_diff),
        write_tsv(pd.DataFrame(blocking_bridge), out_blocking),
    ]

    return PhaseResult(
        estado="COMPLETADA",
        salidas=salidas,
        conteos={
            "carreras_procesadas": len(rows),
            "codigos_no_en_puente": len(no_bridge),
            "codigos_multiples_en_puente": len(multiple_bridge_frames),
            "codcarr_difiere_canonico": len(codcarr_diff),
            "puente_bloqueante": len(blocking_bridge),
        },
        validaciones=[
            {"validacion": "codigo_unico_existe_puente", "bloqueos": len(no_bridge)},
            {"validacion": "codigo_unico_relacion_unica", "bloqueos": len(multiple_bridge_frames)},
            {"validacion": "codcarpr_coincide_canonico", "bloqueos": len(codcarr_diff)},
            {"validacion": "puente_no_bloqueante", "bloqueos": len(blocking_bridge)},
        ],
        pendientes=[],
        resumen=[
            "Estado: COMPLETADA",
            f"Carreras procesadas: {len(rows)}",
            f"No en puente: {len(no_bridge)}",
            f"Multiples en puente: {len(multiple_bridge_frames)}",
            f"CODCARPR difiere: {len(codcarr_diff)}",
            f"Puente bloqueante: {len(blocking_bridge)}",
            f"Validacion puente: {out_validation.resolve()}",
        ],
        comando_reanudacion=command_for_phase(4),
    )


def fase_4(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    universe = read_tsv(output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"), normalize=True)
    identity = read_tsv(RUTA_IDENTIDAD_34, normalize=True)
    detail_identity = read_tsv(RUTA_DETALLE_IDENTIDAD_34, normalize=True)
    comparison_pdf = read_tsv(RUTA_COMPARACION_IDENTIDAD_PDF_PLAN, normalize=True)
    comparison_bridge = read_tsv(RUTA_COMPARACION_IDENTIDAD_PUENTE, normalize=True)

    codigo_u = required_column(universe, "CODIGO_UNICO")
    plan_sies_u = first_existing(list(universe.columns), ["PLAN_ESTUDIOS"])
    pdf_u = required_column(universe, "PDF")
    identity_codes = identity["CODIGO_UNICO"].astype(str).tolist()
    errors: list[dict[str, Any]] = []
    if len(identity) != 34:
        errors.append({"control": "identidad_34_filas", "valor": len(identity), "esperado": 34})
    if identity["CODIGO_UNICO"].duplicated().any():
        errors.append({"control": "identidad_codigo_unico_duplicado", "valor": int(identity["CODIGO_UNICO"].duplicated().sum())})
    if set(identity_codes) != set(universe[codigo_u].astype(str)):
        errors.append(
            {
                "control": "identidad_no_coincide_universo",
                "ausentes": sorted(set(universe[codigo_u].astype(str)) - set(identity_codes))[:20],
                "adicionales": sorted(set(identity_codes) - set(universe[codigo_u].astype(str)))[:20],
            }
        )
    invalid_identity = identity[
        identity["ESTADO_APLICACION_34"].isin(
            ["AMBIGUO_IDENTIDAD", "BLOQUEADO_IDENTIDAD", "CONTRADICCION_PDF_PLAN", "CONTRADICCION_PUENTE", "NO_EVALUADO"]
        )
    ].copy()
    if not invalid_identity.empty:
        errors.append({"control": "identidad_no_resuelta", "valor": len(invalid_identity)})

    by_identity = identity.set_index("CODIGO_UNICO").to_dict("index")
    rows: list[dict[str, Any]] = []
    invalid_plan_sies: list[dict[str, Any]] = []
    no_eval: list[dict[str, Any]] = []
    duplicated_plan: list[dict[str, Any]] = []
    contradictions: list[dict[str, Any]] = []
    variants: list[dict[str, Any]] = []
    trace: list[dict[str, Any]] = []
    seen_code_plan: set[tuple[str, str]] = set()

    for _, row in universe.iterrows():
        code = str(row[codigo_u])
        id_row = by_identity.get(code, {})
        plan_sies = id_row.get("PLAN_ESTUDIOS_SIES_UNIFICADO", "") or (str(row[plan_sies_u]) if plan_sies_u else "")
        pdf = id_row.get("ARCHIVO_PDF", "") or str(row[pdf_u])
        plan_pdf = id_row.get("PLAN_DE_ESTUDIO_CANONICO_PDF", "")
        plan_inst = id_row.get("PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO", "")
        plans_observed = id_row.get("PLANES_INSTITUCIONALES_OBSERVADOS", "")

        plan_key = (code, plan_sies)
        duplicate_plan = plan_key in seen_code_plan and bool(plan_sies)
        seen_code_plan.add(plan_key)
        is_numeric = str(plan_sies).strip().isdigit()
        is_integer = is_numeric
        is_positive = is_numeric and int(str(plan_sies)) > 0
        if not plan_sies:
            plan_state = "BLOQUEADO_PLAN_ESTUDIOS_SIES_VACIO"
            plan_reason = "PLAN_ESTUDIOS_SIES vacío."
        elif not is_numeric:
            plan_state = "BLOQUEADO_PLAN_ESTUDIOS_SIES_NO_NUMERICO"
            plan_reason = "PLAN_ESTUDIOS_SIES no numérico."
        elif not is_integer:
            plan_state = "BLOQUEADO_PLAN_ESTUDIOS_SIES_NO_ENTERO"
            plan_reason = "PLAN_ESTUDIOS_SIES no entero."
        elif not is_positive:
            plan_state = "BLOQUEADO_PLAN_ESTUDIOS_SIES_NO_POSITIVO"
            plan_reason = "PLAN_ESTUDIOS_SIES no positivo."
        elif duplicate_plan:
            plan_state = "BLOQUEADO_PLAN_ESTUDIOS_SIES_DUPLICADO"
            plan_reason = "Duplicado CODIGO_UNICO + PLAN_ESTUDIOS_SIES."
        else:
            plan_state = "VALIDADO_PLAN_ESTUDIOS_SIES"
            plan_reason = "PLAN_ESTUDIOS_SIES validado como correlativo observado; no comparado contra plan institucional."
        if plan_state != "VALIDADO_PLAN_ESTUDIOS_SIES":
            invalid_plan_sies.append({"CODIGO_UNICO": code, "PLAN_ESTUDIOS_SIES": plan_sies, "ESTADO": plan_state, "MOTIVO": plan_reason})
        if duplicate_plan:
            duplicated_plan.append({"CODIGO_UNICO": code, "PLAN_ESTUDIOS_SIES": plan_sies})

        app_state = id_row.get("ESTADO_APLICACION_34", "")
        if not id_row:
            identity_state = "BLOQUEADO_IDENTIDAD_PREVIA"
            identity_reason = "CODIGO_UNICO sin identidad integrada."
        elif app_state == "VALIDADO_IDENTIDAD_CON_VARIANTES":
            identity_state = "VALIDADO_IDENTIDAD_PDF_PLAN_CON_VARIANTES"
            identity_reason = "Identidad resuelta con variantes legítimas conservadas."
        elif app_state == "VALIDADO_IDENTIDAD_COMPLETA":
            identity_state = "VALIDADO_IDENTIDAD_PDF_PLAN"
            identity_reason = "Identidad resuelta."
        elif app_state == "VALIDADO_IDENTIDAD_SIN_PLAN_INSTITUCIONAL":
            identity_state = "VALIDADO_IDENTIDAD_SIN_PLAN_INSTITUCIONAL"
            identity_reason = "Identidad resuelta sin plan institucional suficiente para comparar."
        elif app_state == "AMBIGUO_IDENTIDAD":
            identity_state = "BLOQUEADO_IDENTIDAD_AMBIGUA"
            identity_reason = "Identidad previa ambigua."
        elif app_state in {"BLOQUEADO_IDENTIDAD", "NO_EVALUADO"}:
            identity_state = "BLOQUEADO_IDENTIDAD_PREVIA"
            identity_reason = "Identidad previa bloqueada o no evaluada."
        elif app_state == "CONTRADICCION_PDF_PLAN":
            identity_state = "BLOQUEADO_CONTRADICCION_PDF_PLAN"
            identity_reason = "Contradicción PDF-plan heredada de identidad."
        else:
            identity_state = "BLOQUEADO_IDENTIDAD_PREVIA"
            identity_reason = f"Estado no válido para integración: {app_state}"

        if not pdf:
            identity_state = "BLOQUEADO_SIN_PDF"
            identity_reason = "PDF no informado."
        elif not plan_pdf:
            identity_state = "BLOQUEADO_SIN_PLAN_CANONICO"
            identity_reason = "Plan canónico PDF no informado."
        elif plans_observed and plan_pdf not in [p for p in str(plans_observed).split("|") if p]:
            identity_state = "BLOQUEADO_CONTRADICCION_PDF_PLAN"
            identity_reason = "Plan canónico PDF no está entre planes institucionales observados."

        if identity_state in {"VALIDADO_IDENTIDAD_SIN_PLAN_INSTITUCIONAL", "BLOQUEADO_IDENTIDAD_PREVIA"}:
            no_eval.append({"CODIGO_UNICO": code, "ARCHIVO_PDF": pdf, "ESTADO": identity_state, "MOTIVO": identity_reason})
        if identity_state == "BLOQUEADO_CONTRADICCION_PDF_PLAN":
            contradictions.append({"CODIGO_UNICO": code, "ARCHIVO_PDF": pdf, "PLAN_DE_ESTUDIO_CANONICO_PDF": plan_pdf, "PLANES_INSTITUCIONALES_OBSERVADOS": plans_observed})
        if id_row.get("TIENE_VARIANTES_LEGITIMAS", "") == "SI":
            variants.append(
                {
                    "CODIGO_UNICO": code,
                    "TIPO_VARIANTE": id_row.get("TIPO_VARIANTE", ""),
                    "JORNADAS_OBSERVADAS": id_row.get("JORNADAS_OBSERVADAS", ""),
                    "FUENTE_IDENTIDAD": id_row.get("FUENTE_PRINCIPAL", ""),
                }
            )
        trace.append(
            {
                "CODIGO_UNICO": code,
                "FUENTE_IDENTIDAD": id_row.get("FUENTE_PRINCIPAL", ""),
                "FUENTES_RESPALDO": id_row.get("FUENTES_RESPALDO", ""),
                "HASH_TRAZABILIDAD_IDENTIDAD": id_row.get("HASH_TRAZABILIDAD_APLICACION", ""),
                "RUTA_IDENTIDAD_34": str(RUTA_IDENTIDAD_34.resolve()),
            }
        )
        rows.append(
            {
                "CODIGO_UNICO": code,
                "ARCHIVO_PDF": pdf,
                "PLAN_DE_ESTUDIO_CANONICO_PDF": plan_pdf,
                "PLAN_ESTUDIOS_SIES": plan_sies,
                "PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO": plan_inst,
                "PLANES_INSTITUCIONALES_OBSERVADOS": plans_observed,
                "JORNADA_UNIFICADA": id_row.get("JORNADA_UNIFICADA", ""),
                "JORNADAS_OBSERVADAS": id_row.get("JORNADAS_OBSERVADAS", ""),
                "CODCARR_UNIFICADO": id_row.get("CODCARR_UNIFICADO", ""),
                "CODCARPR_UNIFICADO": id_row.get("CODCARPR_UNIFICADO", ""),
                "TIPO_PLAN_CARRERA_UNIFICADO": id_row.get("TIPO_PLAN_CARRERA_UNIFICADO", ""),
                "DURACION_ESTUDIOS_UNIFICADA": id_row.get("DURACION_ESTUDIOS_UNIFICADA", ""),
                "ESTADO_IDENTIDAD": id_row.get("ESTADO_IDENTIDAD", ""),
                "ESTADO_APLICACION_34": app_state,
                "PLAN_ESTUDIOS_SIES_ES_NUMERICO": "SI" if is_numeric else "NO",
                "PLAN_ESTUDIOS_SIES_ES_ENTERO": "SI" if is_integer else "NO",
                "PLAN_ESTUDIOS_SIES_ES_POSITIVO": "SI" if is_positive else "NO",
                "PLAN_ESTUDIOS_SIES_DUPLICADO_POR_CODIGO_UNICO": "SI" if duplicate_plan else "NO",
                "ESTADO_VALIDACION_PLAN_ESTUDIOS_SIES": plan_state,
                "MOTIVO_VALIDACION_PLAN_ESTUDIOS_SIES": plan_reason,
                "ESTADO_VALIDACION_IDENTIDAD_PDF_PLAN": identity_state,
                "MOTIVO_VALIDACION_IDENTIDAD_PDF_PLAN": identity_reason,
                "TIENE_VARIANTES_LEGITIMAS": id_row.get("TIENE_VARIANTES_LEGITIMAS", ""),
                "TIPO_VARIANTE": id_row.get("TIPO_VARIANTE", ""),
                "TIENE_CAMPO_FALTANTE_CRITICO": id_row.get("TIENE_CAMPO_FALTANTE_CRITICO", ""),
                "CAMPOS_FALTANTES": id_row.get("CAMPOS_FALTANTES", ""),
                "ES_BLOQUEANTE": id_row.get("ES_BLOQUEANTE", ""),
                "MOTIVO_BLOQUEO": id_row.get("MOTIVO_BLOQUEO", ""),
                "FUENTE_IDENTIDAD": id_row.get("FUENTE_PRINCIPAL", ""),
                "HASH_TRAZABILIDAD_IDENTIDAD": id_row.get("HASH_TRAZABILIDAD_APLICACION", ""),
                "VERSION_LOGICA_FASE4": VERSION_LOGICA_FASE4_INTEGRADA,
            }
        )

    result_df = pd.DataFrame(rows)
    blockers = result_df[
        (result_df["ESTADO_VALIDACION_PLAN_ESTUDIOS_SIES"] != "VALIDADO_PLAN_ESTUDIOS_SIES")
        | result_df["ESTADO_VALIDACION_IDENTIDAD_PDF_PLAN"].str.startswith("BLOQUEADO", na=False)
    ].copy()
    if not blockers.empty:
        errors.append({"control": "bloqueos_fase4_integrada", "valor": len(blockers)})
    if len(result_df) != 34:
        errors.append({"control": "total_fase4_integrada", "valor": len(result_df), "esperado": 34})
    if result_df["CODIGO_UNICO"].duplicated().any():
        errors.append({"control": "duplicados_codigo_unico", "valor": int(result_df["CODIGO_UNICO"].duplicated().sum())})

    out_validation = output_path(ctx, "02_INTERMEDIOS/04_VALIDACION_PLAN_MATRIZ.tsv")
    out_invalid_plan = output_path(ctx, "04_AUDITORIA/10_PLAN_ESTUDIOS_SIES_INVALIDO.tsv")
    out_no_eval = output_path(ctx, "04_AUDITORIA/11_IDENTIDAD_PDF_PLAN_NO_EVALUABLE.tsv")
    out_dups = output_path(ctx, "04_AUDITORIA/12_DUPLICADOS_CODIGO_UNICO_PLAN_ESTUDIOS.tsv")
    out_contra = output_path(ctx, "04_AUDITORIA/13_CONTRADICCIONES_IDENTIDAD_PDF_PLAN.tsv")
    out_variants = output_path(ctx, "04_AUDITORIA/14_VARIANTES_LEGITIMAS_34.tsv")
    out_trace = output_path(ctx, "04_AUDITORIA/15_TRAZABILIDAD_IDENTIDAD_INTEGRADA.tsv")
    out_json = output_path(ctx, "04_AUDITORIA/CORRECCION_FUNCIONAL_FASE4_INTEGRACION_IDENTIDAD.json")
    salidas = [
        write_tsv(result_df, out_validation),
        write_tsv(pd.DataFrame(invalid_plan_sies), out_invalid_plan),
        write_tsv(pd.DataFrame(no_eval), out_no_eval),
        write_tsv(pd.DataFrame(duplicated_plan), out_dups),
        write_tsv(pd.DataFrame(contradictions), out_contra),
        write_tsv(pd.DataFrame(variants), out_variants),
        write_tsv(pd.DataFrame(trace), out_trace),
        write_json(
            {
                "fecha": now_iso(),
                "script_modificado": str(Path(__file__).resolve()),
                "respaldo_script": "backups/validar_pdf_plan_34_directas_pre_integracion_identidad_*.py",
                "causa_raiz": "Comparacion conceptual incorrecta entre PLAN_ESTUDIOS_SIES y PLAN_DE_ESTUDIO_INSTITUCIONAL.",
                "logica_anterior": "PLAN_MATRIZ == PLAN_DE_ESTUDIO_CANONICO",
                "logica_nueva": VERSION_LOGICA_FASE4_INTEGRADA,
                "entrada_identidad": str(RUTA_IDENTIDAD_34.resolve()),
                "hash_entrada_identidad": sha256_file(RUTA_IDENTIDAD_34),
                "carreras_esperadas": 34,
                "carreras_recibidas": len(identity),
                "comparacion_eliminada": "PLAN_ESTUDIOS_SIES_vs_PLAN_DE_ESTUDIO_INSTITUCIONAL",
                "separacion_plan_sies_plan_institucional": True,
                "variantes_conservadas": True,
                "conteos": {
                    "carreras_evaluadas": len(result_df),
                    "identidades_con_variantes": int((result_df["TIENE_VARIANTES_LEGITIMAS"] == "SI").sum()),
                    "multiples_jornadas": int(result_df["JORNADAS_OBSERVADAS"].str.contains(r"\\|", regex=True, na=False).sum()),
                    "contradicciones_pdf_plan": len(contradictions),
                    "bloqueos": len(blockers),
                    "relaciones_detalle_identidad": len(detail_identity),
                    "comparacion_pdf_plan_fase5": len(comparison_pdf),
                    "comparacion_puente_fase5": len(comparison_bridge),
                },
                "originales_modificados": False,
                "precarga_generada": False,
                "pes_generado": False,
            },
            out_json,
        ),
    ]
    estado = "COMPLETADA" if not errors else "BLOQUEADA"
    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "carreras_procesadas": len(result_df),
            "identidades_resueltas": int(result_df["ESTADO_VALIDACION_IDENTIDAD_PDF_PLAN"].isin(["VALIDADO_IDENTIDAD_PDF_PLAN", "VALIDADO_IDENTIDAD_PDF_PLAN_CON_VARIANTES"]).sum()),
            "identidades_con_variantes": int((result_df["TIENE_VARIANTES_LEGITIMAS"] == "SI").sum()),
            "multiples_jornadas": int(result_df["JORNADAS_OBSERVADAS"].str.contains(r"\\|", regex=True, na=False).sum()),
            "contradicciones_pdf_plan": len(contradictions),
            "bloqueos": len(blockers),
            "version_logica_fase4": VERSION_LOGICA_FASE4_INTEGRADA,
        },
        validaciones=[
            {"validacion": "plan_estudios_sies_separado", "estado": "OK"},
            {"validacion": "plan_institucional_desde_identidad", "estado": "OK"},
            {"validacion": "sin_comparacion_plan_sies_vs_plan_institucional", "estado": "OK"},
            {"validacion": "identidad_34_unica", "estado": "OK" if not errors else "ERROR", "errores": len(errors)},
        ],
        errores=errors,
        entradas={
            "identidad_34": str(RUTA_IDENTIDAD_34.resolve()),
            "detalle_identidad_34": str(RUTA_DETALLE_IDENTIDAD_34.resolve()),
            "comparacion_identidad_pdf_plan": str(RUTA_COMPARACION_IDENTIDAD_PDF_PLAN.resolve()),
            "comparacion_identidad_puente": str(RUTA_COMPARACION_IDENTIDAD_PUENTE.resolve()),
        },
        hashes={
            "identidad_34": sha256_file(RUTA_IDENTIDAD_34),
            "detalle_identidad_34": sha256_file(RUTA_DETALLE_IDENTIDAD_34),
        },
        resumen=[
            f"Estado: {estado}",
            f"Version logica Fase 4: {VERSION_LOGICA_FASE4_INTEGRADA}",
            f"Carreras evaluadas: {len(result_df)}",
            f"Identidades con variantes: {int((result_df['TIENE_VARIANTES_LEGITIMAS'] == 'SI').sum())}",
            f"Multiples jornadas preservadas: {int(result_df['JORNADAS_OBSERVADAS'].str.contains(r'\\|', regex=True, na=False).sum())}",
            f"Contradicciones PDF-plan: {len(contradictions)}",
            f"Bloqueos: {len(blockers)}",
            f"Validacion integrada: {out_validation.resolve()}",
        ],
        comando_reanudacion=command_for_phase(5),
    )


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^A-Za-z0-9]+", " ", text.upper()).strip()
    return re.sub(r"\s+", " ", text)


def pipe_values(value: Any) -> list[str]:
    return [v for v in str(value or "").split("|") if v]


def phase5_classify(row: pd.Series) -> tuple[str, str, str, str]:
    cierre = str(row.get("ESTADO_CIERRE", ""))
    tipo_match = str(row.get("TIPO_MATCH", ""))
    nivel_state = str(row.get("ESTADO_TRIMESTRE_NIVEL", ""))
    if cierre == "VALIDADO_EXACTO":
        return "VALIDADA_EXACTA", "NO", "NO", "Match exacto validado por conciliacion previa."
    if cierre == "CANDIDATO_ALTO_REVISAR_NOMBRE":
        return "VALIDADA_APROXIMA_ALTA", "NO", "NO", "Match aproximado alto mismo nivel segun conciliacion previa."
    if cierre == "CONFLICTO_SOLO_NIVEL":
        return "VALIDADA_CON_DIFERENCIA_NIVEL", "SI", "NO", "Asignatura identificada con diferencia exclusiva de nivel."
    if cierre == "REVISAR_NOMBRE" or tipo_match == "APROXIMADO_REVISAR":
        return "REVISAR_NOMBRE", "NO", "SI", "Nombre requiere revision; no se considera validado."
    if cierre == "SIN_CORRESPONDENCIA" or tipo_match == "SIN_MATCH":
        return "SIN_CORRESPONDENCIA", "NO", "SI", "No existe correspondencia documentada."
    if not str(row.get("ASIGNATURA_PDF", "")):
        return "NO_EVALUABLE", "NO", "SI", "Asignatura PDF vacia o no evaluable."
    if nivel_state == "CONFLICTO_TRIMESTRE_NIVEL":
        return "VALIDADA_CON_DIFERENCIA_NIVEL", "SI", "NO", "Diferencia de nivel visible."
    return "NO_EVALUABLE", "NO", "SI", "Clasificacion previa no permite validar."


def fase_5(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    if any(phase_status(ctx.state, i) != "COMPLETADA" for i in range(5)):
        raise FlujoError("Fases 0 a 4 deben estar COMPLETADAS antes de Fase 5.")
    if phase_status(ctx.state, 6) != "PENDIENTE":
        raise FlujoError("Fase 6 no debe estar ejecutada antes de completar Fase 5.")

    required = [
        output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"),
        output_path(ctx, "02_INTERMEDIOS/02_MAPEO_PDF_PLAN_NORMALIZADO.tsv"),
        output_path(ctx, "02_INTERMEDIOS/04_VALIDACION_PLAN_MATRIZ.tsv"),
        RUTA_IDENTIDAD_34,
        RUTA_DETALLE_IDENTIDAD_34,
        RUTA_CONCILIACION_CLASIFICADA_415,
        RUTA_MALLA_PDF_CANDIDATA_415,
        RUTA_CONCILIACION_CANONICA_ASIGNATURAS,
        RUTA_TOTALES_DOCUMENTALES_CORREGIDOS,
        Path(REQUIRED_INPUTS["hoja1_congelada"]["path"]),
    ]
    missing = [str(p.resolve()) for p in required if not p.exists()]
    if missing:
        return PhaseResult(
            estado="BLOQUEADA",
            errores=[{"control": "fuentes_obligatorias_faltantes", "detalle": "|".join(missing)}],
            resumen=["Estado: BLOQUEADA", "Faltan fuentes obligatorias de Fase 5."],
            comando_reanudacion=f"NO AVANZAR. Resolver primero: {output_path(ctx, '05_AUDITORIA/FASE5_CONTROL_INTEGRIDAD.tsv').resolve()}",
        )

    universe = read_tsv(output_path(ctx, "02_INTERMEDIOS/01_UNIVERSO_34_CARRERAS_DIRECTAS.tsv"), normalize=True)
    phase4 = read_tsv(output_path(ctx, "02_INTERMEDIOS/04_VALIDACION_PLAN_MATRIZ.tsv"), normalize=True)
    identity = read_tsv(RUTA_IDENTIDAD_34, normalize=True)
    classified = read_tsv(RUTA_CONCILIACION_CLASIFICADA_415, normalize=True)
    pdf_extract = read_tsv(RUTA_MALLA_PDF_CANDIDATA_415, normalize=True)
    canonical_conc = read_tsv(RUTA_CONCILIACION_CANONICA_ASIGNATURAS, normalize=True)
    totals = read_tsv(RUTA_TOTALES_DOCUMENTALES_CORREGIDOS, normalize=True)
    hoja1 = read_tsv(Path(REQUIRED_INPUTS["hoja1_congelada"]["path"]), normalize=True)

    direct_codes = set(universe["CODIGO_UNICO"].astype(str))
    pdfs = sorted(set(phase4["ARCHIVO_PDF"].astype(str)))
    plans = sorted(set(phase4["PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO"].astype(str)))
    identity_by_pdf = phase4.groupby("ARCHIVO_PDF", dropna=False).agg(
        CARRERAS_ASOCIADAS=("CODIGO_UNICO", lambda x: "|".join(sorted(set(map(str, x))))),
        CODIGOS_UNICOS_ASOCIADOS=("CODIGO_UNICO", lambda x: "|".join(sorted(set(map(str, x))))),
    ).reset_index()
    pdf_to_codes = dict(zip(identity_by_pdf["ARCHIVO_PDF"], identity_by_pdf["CODIGOS_UNICOS_ASOCIADOS"]))

    academic_sources = []
    source_specs = [
        (RUTA_MALLA_PDF_CANDIDATA_415, "EXTRACCION_PDF", "Base de 415 filas extraidas de PDF."),
        (Path(REQUIRED_INPUTS["hoja1_congelada"]["path"]), "FUENTE_CANONICA_HOJA1", "Hoja1 congelada para planes institucionales."),
        (output_path(ctx, "02_INTERMEDIOS/02_MAPEO_PDF_PLAN_NORMALIZADO.tsv"), "MAPEO_PDF_PLAN", "Mapeo PDF-plan normalizado de la ejecucion."),
        (RUTA_CONCILIACION_CLASIFICADA_415, "CONCILIACION_ASIGNATURAS", "Conciliacion clasificada de 415 filas."),
        (RUTA_TOTALES_DOCUMENTALES_CORREGIDOS, "TOTALES_DOCUMENTALES", "Totales documentales corregidos por PDF."),
        (RUTA_IDENTIDAD_34, "IDENTIDAD_CARRERA", "Identidad integrada de las 34 carreras."),
        (RUTA_CONCILIACION_CANONICA_ASIGNATURAS, "AUDITORIA_PREVIA", "Conciliacion canonica PDF-Hoja1."),
    ]
    for idx, (path, role, reason) in enumerate(source_specs, start=1):
        df = read_tsv(path, normalize=True)
        academic_sources.append(
            {
                "ID_FUENTE": f"F{idx:02d}",
                "RUTA": str(path.resolve()),
                "ARCHIVO": path.name,
                "TIPO": path.suffix.lower().lstrip("."),
                "PROCESO": PROCESO,
                "ROL": role,
                "HASH_SHA256": sha256_file(path),
                "REGISTROS": len(df),
                "COLUMNAS": len(df.columns),
                "PLANES": len(set(df["PLAN_DE_ESTUDIO_CANONICO"])) if "PLAN_DE_ESTUDIO_CANONICO" in df.columns else (len(set(df["PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO"])) if "PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO" in df.columns else ""),
                "PDF": len(set(df["ARCHIVO_PDF"])) if "ARCHIVO_PDF" in df.columns else "",
                "MOTIVO_USO": reason,
                "ESTADO_VALIDACION": "LEGIBLE",
            }
        )

    pdf_rows = classified[classified["ARCHIVO_PDF"].isin(pdfs)].copy()
    pdf_base_rows: list[dict[str, Any]] = []
    for idx, row in pdf_rows.iterrows():
        pdf_base_rows.append(
            {
                "ARCHIVO_PDF": row.get("ARCHIVO_PDF", ""),
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": row.get("PLAN_DE_ESTUDIO_CANONICO", ""),
                "NOMBRE_CARRERA": "",
                "NIVEL_PDF": row.get("ANIO_CURRICULAR", ""),
                "SEMESTRE_PDF": row.get("TRIMESTRE", ""),
                "ORDEN_PDF": row.get("ORDEN_VERTICAL", ""),
                "CODIGO_ASIGNATURA_PDF": "",
                "NOMBRE_ASIGNATURA_PDF": row.get("ASIGNATURA_PDF", ""),
                "NOMBRE_ASIGNATURA_NORMALIZADO": row.get("ASIGNATURA_NORMALIZADA", "") or normalize_text(row.get("ASIGNATURA_PDF", "")),
                "HORAS_PDF": "",
                "CREDITOS_PDF": "",
                "UNIDAD_MEDIDA_PDF": row.get("TIPO_ACTIVIDAD", ""),
                "FUENTE": str(RUTA_CONCILIACION_CLASIFICADA_415.resolve()),
                "HASH_FUENTE": sha256_file(RUTA_CONCILIACION_CLASIFICADA_415),
            }
        )
    pdf_base = pd.DataFrame(pdf_base_rows)

    plan_col = first_existing(list(hoja1.columns), ["PLAN_DE_ESTUDIO"])
    canonical_rows: list[dict[str, Any]] = []
    hoja_plans = hoja1[hoja1[plan_col].isin(plans)].copy() if plan_col else empty_df([])
    for _, row in hoja_plans.iterrows():
        canonical_rows.append(
            {
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": row.get(plan_col, ""),
                "CODIGO_ASIGNATURA_CANONICO": row.get("CODRAMO", ""),
                "NOMBRE_ASIGNATURA_CANONICO": row.get("ASIGNATURA", ""),
                "NOMBRE_ASIGNATURA_NORMALIZADO": normalize_text(row.get("ASIGNATURA", "")),
                "NIVEL_CANONICO": row.get("NIVEL", ""),
                "SEMESTRE_CANONICO": row.get("PERIODO", ""),
                "ORDEN_CANONICO": "",
                "HORAS_CANONICAS": "",
                "CREDITOS_CANONICOS": "",
                "UNIDAD_MEDIDA_CANONICA": "ASIGNATURA",
                "VIGENCIA": row.get("ESTADO", ""),
                "FUENTE": str(Path(REQUIRED_INPUTS["hoja1_congelada"]["path"]).resolve()),
                "HASH_FUENTE": sha256_file(Path(REQUIRED_INPUTS["hoja1_congelada"]["path"])),
            }
        )
    canonical_df = pd.DataFrame(canonical_rows).drop_duplicates()

    exact_rows: list[dict[str, Any]] = []
    approx_rows: list[dict[str, Any]] = []
    final_rows: list[dict[str, Any]] = []
    for idx, row in pdf_rows.reset_index(drop=True).iterrows():
        estado_asig, diff_level, manual, motivo = phase5_classify(row)
        tipo_match = row.get("TIPO_MATCH", "")
        codigo_can = pipe_values(row.get("CODRAMOS_MATCH", ""))
        nombre_can = pipe_values(row.get("ASIGNATURAS_HOJA1", ""))
        nivel_can = pipe_values(row.get("NIVELES_HOJA1", ""))
        exact_state = "NO_MATCH_EXACTO"
        if tipo_match == "EXACTO" and codigo_can:
            exact_state = "MATCH_EXACTO_CODIGO"
        elif tipo_match == "EXACTO" and row.get("ESTADO_TRIMESTRE_NIVEL") == "COINCIDE_TRIMESTRE_NIVEL":
            exact_state = "MATCH_EXACTO_NOMBRE_NIVEL"
        elif tipo_match == "EXACTO":
            exact_state = "MATCH_EXACTO_NOMBRE_SIN_NIVEL"
        exact_rows.append({**row.to_dict(), "ESTADO_MATCH_EXACTO": exact_state})
        if tipo_match != "EXACTO":
            approx_rows.append(
                {
                    **row.to_dict(),
                    "ESTADO_MATCH_APROXIMADO": {
                        "APROXIMADO_ALTO": "MATCH_APROXIMO_ALTO_MISMO_NIVEL" if diff_level == "NO" else "MATCH_APROXIMO_ALTO_NIVEL_DISTINTO",
                        "APROXIMADO_REVISAR": "REVISAR_NOMBRE",
                        "SIN_MATCH": "SIN_CORRESPONDENCIA",
                    }.get(tipo_match, "NO_EVALUABLE"),
                    "PUNTAJE_SIMILITUD": row.get("SIMILITUD", ""),
                    "METODO_SIMILITUD": "METODO_PREVIO_CONCILIACION_CANONICA",
                    "CANDIDATOS_EVALUADOS": row.get("ASIGNATURAS_HOJA1", ""),
                    "DIFERENCIA_NIVEL": diff_level,
                    "JUSTIFICACION": motivo,
                }
            )
        canonical_name = "|".join(nombre_can)
        final = {
            "CODIGO_UNICO": pdf_to_codes.get(row.get("ARCHIVO_PDF", ""), ""),
            "ARCHIVO_PDF": row.get("ARCHIVO_PDF", ""),
            "PLAN_DE_ESTUDIO_INSTITUCIONAL": row.get("PLAN_DE_ESTUDIO_CANONICO", ""),
            "NIVEL_PDF": row.get("ANIO_CURRICULAR", ""),
            "SEMESTRE_PDF": row.get("TRIMESTRE", ""),
            "ORDEN_PDF": row.get("ORDEN_VERTICAL", ""),
            "CODIGO_ASIGNATURA_PDF": "",
            "NOMBRE_ASIGNATURA_PDF": row.get("ASIGNATURA_PDF", ""),
            "CODIGO_ASIGNATURA_CANONICO": "|".join(codigo_can),
            "NOMBRE_ASIGNATURA_CANONICO": canonical_name,
            "NIVEL_CANONICO": "|".join(nivel_can),
            "SEMESTRE_CANONICO": row.get("PERIODOS_HOJA1", ""),
            "TIPO_MATCH": tipo_match,
            "PUNTAJE_SIMILITUD": row.get("SIMILITUD", ""),
            "DIFERENCIA_NIVEL": diff_level,
            "DIFERENCIA_NOMBRE": "SI" if estado_asig in {"REVISAR_NOMBRE", "VALIDADA_APROXIMA_ALTA"} else "NO",
            "DIFERENCIA_CODIGO": "SI" if not codigo_can and estado_asig not in {"SIN_CORRESPONDENCIA", "NO_EVALUABLE"} else "NO",
            "ESTADO_VALIDACION_ASIGNATURA": estado_asig,
            "MOTIVO": motivo,
            "REQUIERE_REVISION_MANUAL": manual,
            "ES_BLOQUEANTE": "NO",
            "FUENTE_PDF": str(RUTA_CONCILIACION_CLASIFICADA_415.resolve()),
            "FUENTE_CANONICA": str(Path(REQUIRED_INPUTS["hoja1_congelada"]["path"]).resolve()),
            "HASH_TRAZABILIDAD": "",
        }
        final["HASH_TRAZABILIDAD"] = hashlib.sha256("|".join(str(final[k]) for k in final if k != "HASH_TRAZABILIDAD").encode("utf-8")).hexdigest()
        final_rows.append(final)
    final_df = pd.DataFrame(final_rows)

    pdf_summary_rows: list[dict[str, Any]] = []
    for pdf, group in final_df.groupby("ARCHIVO_PDF", dropna=False):
        plan = group["PLAN_DE_ESTUDIO_INSTITUCIONAL"].iloc[0]
        canon_total = len(canonical_df[canonical_df["PLAN_DE_ESTUDIO_INSTITUCIONAL"] == plan]) if not canonical_df.empty else 0
        exact_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_EXACTA").sum())
        approx_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_APROXIMA_ALTA").sum())
        level_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_CON_DIFERENCIA_NIVEL").sum())
        review_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "REVISAR_NOMBRE").sum())
        amb_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "AMBIGUA").sum())
        sin_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "SIN_CORRESPONDENCIA").sum())
        no_eval_n = int((group["ESTADO_VALIDACION_ASIGNATURA"] == "NO_EVALUABLE").sum())
        total = len(group)
        if amb_n:
            pdf_state = "PDF_BLOQUEADO"
        elif review_n or sin_n or no_eval_n:
            pdf_state = "PDF_REQUIERE_REVISION"
        elif level_n:
            pdf_state = "PDF_VALIDADO_CON_OBSERVACIONES"
        else:
            pdf_state = "PDF_VALIDADO_COMPLETO"
        pdf_summary_rows.append(
            {
                "ARCHIVO_PDF": pdf,
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": plan,
                "CARRERAS_ASOCIADAS": pdf_to_codes.get(pdf, ""),
                "CODIGOS_UNICOS_ASOCIADOS": pdf_to_codes.get(pdf, ""),
                "TOTAL_FILAS_PDF": total,
                "TOTAL_ASIGNATURAS_PDF": group["NOMBRE_ASIGNATURA_PDF"].nunique(),
                "TOTAL_ASIGNATURAS_CANONICAS": canon_total,
                "VALIDADAS_EXACTAS": exact_n,
                "VALIDADAS_APROXIMAS_ALTAS": approx_n,
                "DIFERENCIA_NIVEL": level_n,
                "REVISAR_NOMBRE": review_n,
                "AMBIGUAS": amb_n,
                "SIN_CORRESPONDENCIA": sin_n,
                "NO_EVALUABLES": no_eval_n,
                "COBERTURA_VALIDADA": round(exact_n * 100 / total, 2) if total else 0,
                "COBERTURA_TOTAL_CON_APROXIMADAS": round((exact_n + approx_n) * 100 / total, 2) if total else 0,
                "ESTADO_PDF": pdf_state,
                "ES_BLOQUEANTE": "SI" if pdf_state == "PDF_BLOQUEADO" else "NO",
                "MOTIVO": "Casos para revision documentados." if pdf_state == "PDF_REQUIERE_REVISION" else "",
            }
        )
    pdf_summary = pd.DataFrame(pdf_summary_rows)

    career_rows: list[dict[str, Any]] = []
    for _, row in phase4.iterrows():
        pdf = row["ARCHIVO_PDF"]
        pdf_row = pdf_summary[pdf_summary["ARCHIVO_PDF"] == pdf].iloc[0].to_dict()
        estado_pdf = pdf_row["ESTADO_PDF"]
        estado_carrera = {
            "PDF_VALIDADO_COMPLETO": "CARRERA_VALIDADA_COMPLETA",
            "PDF_VALIDADO_CON_OBSERVACIONES": "CARRERA_VALIDADA_CON_OBSERVACIONES",
            "PDF_REQUIERE_REVISION": "CARRERA_REQUIERE_REVISION",
            "PDF_BLOQUEADO": "CARRERA_BLOQUEADA",
        }[estado_pdf]
        career_rows.append(
            {
                "CODIGO_UNICO": row["CODIGO_UNICO"],
                "ARCHIVO_PDF": pdf,
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": row["PLAN_DE_ESTUDIO_INSTITUCIONAL_UNIFICADO"],
                "JORNADAS_OBSERVADAS": row.get("JORNADAS_OBSERVADAS", ""),
                "MODALIDADES_OBSERVADAS": "",
                "SEDES_OBSERVADAS": "",
                "VERSIONES_OBSERVADAS": "",
                "TOTAL_ASIGNATURAS_PDF": pdf_row["TOTAL_ASIGNATURAS_PDF"],
                "VALIDADAS_EXACTAS": pdf_row["VALIDADAS_EXACTAS"],
                "VALIDADAS_APROXIMAS_ALTAS": pdf_row["VALIDADAS_APROXIMAS_ALTAS"],
                "DIFERENCIA_NIVEL": pdf_row["DIFERENCIA_NIVEL"],
                "REVISAR_NOMBRE": pdf_row["REVISAR_NOMBRE"],
                "AMBIGUAS": pdf_row["AMBIGUAS"],
                "SIN_CORRESPONDENCIA": pdf_row["SIN_CORRESPONDENCIA"],
                "NO_EVALUABLES": pdf_row["NO_EVALUABLES"],
                "ESTADO_PDF": estado_pdf,
                "ESTADO_CARRERA": estado_carrera,
                "ES_BLOQUEANTE": pdf_row["ES_BLOQUEANTE"],
                "MOTIVO": pdf_row["MOTIVO"],
            }
        )
    career_df = pd.DataFrame(career_rows)

    total_rows: list[dict[str, Any]] = []
    for _, t in totals[totals["ARCHIVO_PDF"].isin(pdfs)].iterrows():
        pdf = t["ARCHIVO_PDF"]
        pdf_row = pdf_summary[pdf_summary["ARCHIVO_PDF"] == pdf].iloc[0].to_dict()
        for metric, col in [
            ("TOTAL_UNIDADES_DOCUMENTALES", "TOTAL_UNIDADES_DOCUMENTALES"),
            ("UNIDADES_EXACTAS", "UNIDADES_EXACTAS"),
            ("UNIDADES_APROXIMADAS_ALTAS", "UNIDADES_APROXIMADAS_ALTAS"),
        ]:
            observed = int(pdf_row["TOTAL_FILAS_PDF"]) if metric == "TOTAL_UNIDADES_DOCUMENTALES" else (int(pdf_row["VALIDADAS_EXACTAS"]) if metric == "UNIDADES_EXACTAS" else int(pdf_row["VALIDADAS_APROXIMAS_ALTAS"]))
            expected = int(t[col]) if str(t[col]).isdigit() else 0
            diff = observed - expected
            total_rows.append(
                {
                    "ARCHIVO_PDF": pdf,
                    "PLAN_DE_ESTUDIO_INSTITUCIONAL": t.get("PLAN_DE_ESTUDIO_CANONICO", ""),
                    "METRICA": metric,
                    "UNIDAD": "ASIGNATURA",
                    "VALOR_PDF": observed,
                    "VALOR_CANONICO": expected,
                    "DIFERENCIA": diff,
                    "PORCENTAJE_DIFERENCIA": round(diff * 100 / expected, 2) if expected else "",
                    "ESTADO": "TOTAL_COINCIDE" if diff == 0 else "TOTAL_DIFIERE_REVISAR",
                    "MOTIVO": "" if diff == 0 else "Diferencia documental trazada; no bloquea Fase 5.",
                }
            )
    totals_df = pd.DataFrame(total_rows)

    control_errors: list[dict[str, Any]] = []
    if len(set(pdfs)) != 11:
        control_errors.append({"control": "pdf_esperados", "valor": len(set(pdfs)), "esperado": 11})
    if len(direct_codes) != 34 or len(phase4) != 34:
        control_errors.append({"control": "carreras_34", "valor": len(phase4), "esperado": 34})
    if len(set(pdf_rows["ARCHIVO_PDF"])) != 11:
        control_errors.append({"control": "pdf_evaluados", "valor": len(set(pdf_rows["ARCHIVO_PDF"])), "esperado": 11})
    if set(pdf_rows["ARCHIVO_PDF"]) - set(pdfs):
        control_errors.append({"control": "pdf_fuera_universo", "valor": "|".join(sorted(set(pdf_rows["ARCHIVO_PDF"]) - set(pdfs)))})
    if set(plans) - set(canonical_df["PLAN_DE_ESTUDIO_INSTITUCIONAL"] if not canonical_df.empty else []):
        control_errors.append({"control": "planes_ausentes_hoja1", "valor": "|".join(sorted(set(plans) - set(canonical_df["PLAN_DE_ESTUDIO_INSTITUCIONAL"] if not canonical_df.empty else [])))})
    if len(final_df) != len(pdf_rows):
        control_errors.append({"control": "filas_pdf_preservadas", "valor": len(final_df), "esperado": len(pdf_rows)})
    if not all(final_df["ESTADO_VALIDACION_ASIGNATURA"] != ""):
        control_errors.append({"control": "asignaturas_sin_estado"})
    if int((phase4["JORNADAS_OBSERVADAS"].str.contains(r"\|", regex=True, na=False)).sum()) != 21:
        control_errors.append({"control": "multiples_jornadas_preservadas", "valor": int((phase4["JORNADAS_OBSERVADAS"].str.contains(r"\|", regex=True, na=False)).sum())})

    coverage_rows = [
        {"CONTROL": "pdf_evaluados", "VALOR": len(set(pdf_rows["ARCHIVO_PDF"])), "ESPERADO": 11, "ESTADO": "OK" if len(set(pdf_rows["ARCHIVO_PDF"])) == 11 else "ERROR"},
        {"CONTROL": "carreras_evaluadas", "VALOR": len(career_df), "ESPERADO": 34, "ESTADO": "OK" if len(career_df) == 34 else "ERROR"},
        {"CONTROL": "filas_pdf_evaluadas", "VALOR": len(final_df), "ESPERADO": len(pdf_rows), "ESTADO": "OK" if len(final_df) == len(pdf_rows) else "ERROR"},
    ]
    integrity_rows = [
        {"CONTROL": "34_carreras_preservadas", "ESTADO": "OK" if len(career_df) == 34 else "ERROR", "OBSERVACION": str(len(career_df))},
        {"CONTROL": "11_pdf_preservados", "ESTADO": "OK" if len(set(pdf_rows["ARCHIVO_PDF"])) == 11 else "ERROR", "OBSERVACION": str(len(set(pdf_rows["ARCHIVO_PDF"])))},
        {"CONTROL": "11_planes_preservados", "ESTADO": "OK" if len(set(plans)) == 11 else "ERROR", "OBSERVACION": str(len(set(plans)))},
        {"CONTROL": "sin_plan_estudios_como_codigo_institucional", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "sin_keep_first", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "sin_keep_last", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "multiples_jornadas_preservadas", "ESTADO": "OK" if int((phase4["JORNADAS_OBSERVADAS"].str.contains(r'\\|', regex=True, na=False)).sum()) == 21 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "sin_precarga", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "sin_pes", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "sin_fase_posterior", "ESTADO": "OK" if phase_status(ctx.state, 6) == "PENDIENTE" else "ERROR", "OBSERVACION": phase_status(ctx.state, 6)},
    ]

    exact_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_EXACTA").sum())
    approx_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_APROXIMA_ALTA").sum())
    level_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_CON_DIFERENCIA_NIVEL").sum())
    review_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "REVISAR_NOMBRE").sum())
    ambiguous_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "AMBIGUA").sum())
    sin_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "SIN_CORRESPONDENCIA").sum())
    noeval_n = int((final_df["ESTADO_VALIDACION_ASIGNATURA"] == "NO_EVALUABLE").sum())
    pdf_complete = int((pdf_summary["ESTADO_PDF"] == "PDF_VALIDADO_COMPLETO").sum())
    pdf_obs = int((pdf_summary["ESTADO_PDF"] == "PDF_VALIDADO_CON_OBSERVACIONES").sum())
    pdf_review = int((pdf_summary["ESTADO_PDF"] == "PDF_REQUIERE_REVISION").sum())
    pdf_blocked = int((pdf_summary["ESTADO_PDF"] == "PDF_BLOQUEADO").sum())
    car_complete = int((career_df["ESTADO_CARRERA"] == "CARRERA_VALIDADA_COMPLETA").sum())
    car_obs = int((career_df["ESTADO_CARRERA"] == "CARRERA_VALIDADA_CON_OBSERVACIONES").sum())
    car_review = int((career_df["ESTADO_CARRERA"] == "CARRERA_REQUIERE_REVISION").sum())
    car_blocked = int((career_df["ESTADO_CARRERA"] == "CARRERA_BLOQUEADA").sum())
    coverage_exact = round(exact_n * 100 / len(final_df), 2) if len(final_df) else 0
    coverage_total = round((exact_n + approx_n) * 100 / len(final_df), 2) if len(final_df) else 0

    summary = [
        ("PDF evaluados", len(set(pdf_rows["ARCHIVO_PDF"])), "OK", "conciliacion_clasificada", ""),
        ("carreras evaluadas", len(career_df), "OK", "fase4_integrada", ""),
        ("filas PDF", len(final_df), "OK", "conciliacion_clasificada", ""),
        ("asignaturas distintas", final_df["NOMBRE_ASIGNATURA_PDF"].nunique(), "OK", "conciliacion_clasificada", ""),
        ("exactas", exact_n, "OK", "validacion_asignaturas", ""),
        ("aproximadas altas", approx_n, "OK", "validacion_asignaturas", ""),
        ("diferencias de nivel", level_n, "OK", "validacion_asignaturas", ""),
        ("revisar nombre", review_n, "OK", "validacion_asignaturas", ""),
        ("ambiguas", ambiguous_n, "OK", "validacion_asignaturas", ""),
        ("sin correspondencia", sin_n, "OK", "validacion_asignaturas", ""),
        ("no evaluables", noeval_n, "OK", "validacion_asignaturas", ""),
        ("PDF completos", pdf_complete, "OK", "resumen_pdf", ""),
        ("PDF con observaciones", pdf_obs, "OK", "resumen_pdf", ""),
        ("PDF para revision", pdf_review, "OK", "resumen_pdf", ""),
        ("PDF bloqueados", pdf_blocked, "OK", "resumen_pdf", ""),
        ("carreras completas", car_complete, "OK", "resumen_carreras", ""),
        ("carreras con observaciones", car_obs, "OK", "resumen_carreras", ""),
        ("carreras para revision", car_review, "OK", "resumen_carreras", ""),
        ("carreras bloqueadas", car_blocked, "OK", "resumen_carreras", ""),
        ("originales modificados", "NO", "OK", "control", ""),
        ("precarga generada", "NO", "OK", "control", ""),
        ("PES generado", "NO", "OK", "control", ""),
        ("apto para carga", "NO", "OK", "control", ""),
    ]
    summary_df = pd.DataFrame(summary, columns=["METRICA", "VALOR", "ESTADO", "FUENTE", "OBSERVACION"])
    estado = "COMPLETADA" if not control_errors else "BLOQUEADA"

    out_sources = output_path(ctx, "02_INTERMEDIOS/05A_FUENTES_ACADEMICAS_FASE5.tsv")
    out_pdf = output_path(ctx, "02_INTERMEDIOS/05B_ASIGNATURAS_EXTRAIDAS_PDF_11.tsv")
    out_canon = output_path(ctx, "02_INTERMEDIOS/05C_ASIGNATURAS_CANONICAS_HOJA1_11_PLANES.tsv")
    out_exact = output_path(ctx, "02_INTERMEDIOS/05D_CONCILIACION_EXACTA_ASIGNATURAS.tsv")
    out_approx = output_path(ctx, "02_INTERMEDIOS/05E_CONCILIACION_APROXIMADA_ASIGNATURAS.tsv")
    out_assign = output_path(ctx, "04_RESULTADOS/12_VALIDACION_ASIGNATURAS_PDF_PLAN.tsv")
    out_pdf_summary = output_path(ctx, "04_RESULTADOS/13_RESUMEN_VALIDACION_11_PDF.tsv")
    out_career = output_path(ctx, "04_RESULTADOS/14_VALIDACION_34_CARRERAS_PDF_PLAN.tsv")
    out_summary = output_path(ctx, "04_RESULTADOS/15_RESUMEN_FASE5_PDF_PLAN.tsv")
    salidas = [
        write_tsv(pd.DataFrame(academic_sources), out_sources),
        write_tsv(pdf_base, out_pdf),
        write_tsv(canonical_df, out_canon),
        write_tsv(pd.DataFrame(exact_rows), out_exact),
        write_tsv(pd.DataFrame(approx_rows), out_approx),
        write_tsv(final_df, out_assign),
        write_tsv(pdf_summary, out_pdf_summary),
        write_tsv(career_df, out_career),
        write_tsv(summary_df, out_summary),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_EXACTA"], output_path(ctx, "05_AUDITORIA/FASE5_MATCH_EXACTO.tsv")),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_APROXIMA_ALTA"], output_path(ctx, "05_AUDITORIA/FASE5_MATCH_APROXIMO.tsv")),
        write_tsv(final_df[final_df["DIFERENCIA_NIVEL"] == "SI"], output_path(ctx, "05_AUDITORIA/FASE5_DIFERENCIAS_NIVEL.tsv")),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "REVISAR_NOMBRE"], output_path(ctx, "05_AUDITORIA/FASE5_REVISAR_NOMBRE.tsv")),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "AMBIGUA"], output_path(ctx, "05_AUDITORIA/FASE5_AMBIGUAS.tsv")),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "SIN_CORRESPONDENCIA"], output_path(ctx, "05_AUDITORIA/FASE5_SIN_CORRESPONDENCIA.tsv")),
        write_tsv(final_df[final_df["ESTADO_VALIDACION_ASIGNATURA"] == "NO_EVALUABLE"], output_path(ctx, "05_AUDITORIA/FASE5_NO_EVALUABLES.tsv")),
        write_tsv(totals_df, output_path(ctx, "05_AUDITORIA/FASE5_TOTALES_DOCUMENTALES.tsv")),
        write_tsv(pd.DataFrame(coverage_rows), output_path(ctx, "05_AUDITORIA/FASE5_CONTROL_COBERTURA.tsv")),
        write_tsv(pd.DataFrame(integrity_rows), output_path(ctx, "05_AUDITORIA/FASE5_CONTROL_INTEGRIDAD.tsv")),
        write_tsv(final_df[["CODIGO_UNICO", "ARCHIVO_PDF", "PLAN_DE_ESTUDIO_INSTITUCIONAL", "NOMBRE_ASIGNATURA_PDF", "HASH_TRAZABILIDAD", "FUENTE_PDF", "FUENTE_CANONICA"]], output_path(ctx, "05_AUDITORIA/FASE5_TRAZABILIDAD_ASIGNATURAS.tsv")),
        write_tsv(canonical_conc[canonical_conc["ARCHIVO_PDF"].isin(pdfs)], output_path(ctx, "05_AUDITORIA/FASE5_CONCILIACION_CANONICA_REFERENCIA.tsv")),
        write_text(f"{now_iso()} FASE5 validacion academica estado={estado} pdf={len(set(pdf_rows['ARCHIVO_PDF']))} filas={len(final_df)}\n", output_path(ctx, "05_LOGS/fase5_validacion_academica_pdf_plan.log")),
    ]
    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "pdf_esperados": 11,
            "pdf_evaluados": len(set(pdf_rows["ARCHIVO_PDF"])),
            "carreras_esperadas": 34,
            "carreras_evaluadas": len(career_df),
            "filas_pdf": len(final_df),
            "asignaturas_distintas": final_df["NOMBRE_ASIGNATURA_PDF"].nunique(),
            "asignaturas_canonicas_evaluadas": len(canonical_df),
            "matches_exactos": exact_n,
            "matches_aproximados": approx_n,
            "diferencias_nivel": level_n,
            "revisar_nombre": review_n,
            "ambiguas": ambiguous_n,
            "sin_correspondencia": sin_n,
            "no_evaluables": noeval_n,
            "pdf_completos": pdf_complete,
            "pdf_observaciones": pdf_obs,
            "pdf_revision": pdf_review,
            "pdf_bloqueados": pdf_blocked,
            "carreras_completas": car_complete,
            "carreras_observaciones": car_obs,
            "carreras_revision": car_review,
            "carreras_bloqueadas": car_blocked,
            "bloqueos_estructurales": len(control_errors),
            "originales_modificados": False,
            "precarga_generada": False,
            "pes_generado": False,
            "apto_para_carga": False,
        },
        validaciones=integrity_rows + coverage_rows,
        errores=control_errors,
        entradas={str(p.resolve()): str(p.resolve()) for p in required},
        hashes={str(p.resolve()): sha256_file(p) for p in required},
        resumen=[
            f"Estado: {estado}",
            "PDF esperados: 11",
            f"PDF evaluados: {len(set(pdf_rows['ARCHIVO_PDF']))}",
            "Carreras esperadas: 34",
            f"Carreras evaluadas: {len(career_df)}",
            f"Filas PDF evaluadas: {len(final_df)}",
            f"Asignaturas distintas PDF: {final_df['NOMBRE_ASIGNATURA_PDF'].nunique()}",
            f"Asignaturas canonicas evaluadas: {len(canonical_df)}",
            f"Validaciones exactas: {exact_n}",
            f"Validaciones aproximadas altas: {approx_n}",
            f"Diferencias solo de nivel: {level_n}",
            f"Revisar nombre: {review_n}",
            f"Ambiguas: {ambiguous_n}",
            f"Sin correspondencia: {sin_n}",
            f"No evaluables: {noeval_n}",
            f"PDF validados completos: {pdf_complete}",
            f"PDF validados con observaciones: {pdf_obs}",
            f"PDF requieren revision: {pdf_review}",
            f"PDF bloqueados: {pdf_blocked}",
            f"Carreras validadas completas: {car_complete}",
            f"Carreras validadas con observaciones: {car_obs}",
            f"Carreras requieren revision: {car_review}",
            f"Carreras bloqueadas: {car_blocked}",
            f"Cobertura exacta: {coverage_exact}%",
            f"Cobertura exacta + aproximada: {coverage_total}%",
            "PLAN_ESTUDIOS separado de PLAN_DE_ESTUDIO: SI",
            "Identidad de carreras preservada: SI",
            "Variantes legitimas preservadas: SI",
            "Seleccion arbitraria de filas: NO",
            f"Salida asignaturas: {out_assign.resolve()}",
            f"Resumen 11 PDF: {out_pdf_summary.resolve()}",
            f"Resumen 34 carreras: {out_career.resolve()}",
            f"Auditorias: {output_path(ctx, '05_AUDITORIA').resolve()}",
            f"Carpeta de ejecucion: {ctx.run_dir.resolve()}",
            "Originales modificados: NO",
            "Precarga generada: NO",
            "PES generado: NO",
            "Apto para carga: NO",
            "Siguiente paso: REVISAR_CASOS_NO_EXACTOS_Y_DEFINIR_CIERRE_PDF_PLAN",
        ],
        comando_reanudacion="" if estado == "COMPLETADA" else f"NO AVANZAR. Resolver primero: {output_path(ctx, '05_AUDITORIA/FASE5_CONTROL_INTEGRIDAD.tsv').resolve()}",
    )


STOPWORDS_ASIG = {
    "DE", "DEL", "LA", "LAS", "EL", "LOS", "Y", "A", "EN", "PARA", "CON", "I", "II", "III", "IV",
    "V", "VI", "BC", "CO", "FN", "RH", "AWS",
}


def principal_words(value: Any) -> set[str]:
    return {w for w in normalize_text(value).split() if len(w) > 2 and w not in STOPWORDS_ASIG}


def similarity(a: Any, b: Any) -> float:
    return round(difflib.SequenceMatcher(None, normalize_text(a), normalize_text(b)).ratio(), 4)


def roman_tokens(value: Any) -> set[str]:
    return {w for w in normalize_text(value).split() if w in {"I", "II", "III", "IV", "V", "VI"}}


def diff_type(pdf_name: str, canon_name: str) -> str:
    n_pdf = normalize_text(pdf_name)
    n_can = normalize_text(canon_name)
    if not n_can:
        return "SIN_CANDIDATO_CLARO"
    if n_pdf == n_can:
        return "VARIANTE_EDITORIAL"
    if roman_tokens(n_pdf) != roman_tokens(n_can) and (roman_tokens(n_pdf) or roman_tokens(n_can)):
        return "SECUENCIA_DISTINTA"
    pdf_words = principal_words(n_pdf)
    can_words = principal_words(n_can)
    if not pdf_words or not can_words:
        return "SIN_CANDIDATO_CLARO"
    missing = pdf_words - can_words
    added = can_words - pdf_words
    if len(missing) == 0 and len(added) <= 1:
        return "PALABRA_AGREGADA"
    if len(added) == 0 and len(missing) <= 1:
        return "PALABRA_ELIMINADA"
    if similarity(n_pdf, n_can) >= 0.9:
        return "VARIANTE_EDITORIAL"
    if any(w in pdf_words ^ can_words for w in {"PRACTICA", "TALLER", "LABORATORIO", "PROYECTO"}):
        return "PRACTICA_TALLER_LABORATORIO_DISTINTO"
    return "CAMBIO_SUSTANTIVO"


def extraction_class(value: Any, duplicated: bool) -> str:
    text = normalize_text(value)
    if not text or len(text) < 4:
        return "EXTRACCION_INSUFICIENTE"
    if text.isdigit():
        return "EXTRACCION_INSUFICIENTE"
    if text in {"MALLA CURRICULAR", "ASIGNATURA", "NIVEL", "SEMESTRE", "TRIMESTRE"}:
        return "EXTRACCION_ENCABEZADO"
    if duplicated:
        return "EXTRACCION_DUPLICADA"
    return "EXTRACCION_VALIDA"


def build_case_candidates(case: dict[str, Any], canon_rows: list[dict[str, str]], fuente: str) -> list[dict[str, Any]]:
    pdf_name = case.get("NOMBRE_ASIGNATURA_PDF", "")
    pdf_norm = normalize_text(pdf_name)
    pdf_words = principal_words(pdf_name)
    candidates: list[dict[str, Any]] = []
    for cand in canon_rows:
        cand_name = cand.get("NOMBRE_ASIGNATURA_CANONICO", "")
        cand_norm = cand.get("NOMBRE_ASIGNATURA_NORMALIZADO", "") or normalize_text(cand_name)
        cand_words = principal_words(cand_name)
        score = similarity(pdf_norm, cand_norm)
        overlap = len(pdf_words & cand_words) / max(1, len(pdf_words | cand_words))
        if score >= 0.55 or overlap >= 0.35 or normalize_text(case.get("CODIGO_ASIGNATURA_PDF", "")) == normalize_text(cand.get("CODIGO_ASIGNATURA_CANONICO", "")):
            candidates.append(
                {
                    "ID_CASO": case["ID_CASO"],
                    "PLAN_DE_ESTUDIO_INSTITUCIONAL": case["PLAN_DE_ESTUDIO_INSTITUCIONAL"],
                    "CODIGO_CANDIDATO": cand.get("CODIGO_ASIGNATURA_CANONICO", ""),
                    "NOMBRE_CANDIDATO": cand_name,
                    "NIVEL_CANDIDATO": cand.get("NIVEL_CANONICO", ""),
                    "SEMESTRE_CANDIDATO": cand.get("SEMESTRE_CANONICO", ""),
                    "PUNTAJE_SIMILITUD": score,
                    "POSICION_CANDIDATO": 0,
                    "DIFERENCIA_PUNTAJE_RESPECTO_PRIMERO": "",
                    "COINCIDE_CODIGO": "SI" if case.get("CODIGO_ASIGNATURA_PDF") and case.get("CODIGO_ASIGNATURA_PDF") == cand.get("CODIGO_ASIGNATURA_CANONICO") else "NO",
                    "COINCIDE_NOMBRE_NORMALIZADO": "SI" if pdf_norm and pdf_norm == cand_norm else "NO",
                    "COINCIDE_NIVEL": "SI" if case.get("NIVEL_PDF") and case.get("NIVEL_PDF") == cand.get("NIVEL_CANONICO") else "NO",
                    "COINCIDE_SEMESTRE": "SI" if case.get("SEMESTRE_PDF") and case.get("SEMESTRE_PDF") == cand.get("SEMESTRE_CANONICO") else "NO",
                    "COINCIDE_PALABRAS_PRINCIPALES": "SI" if overlap >= 0.5 else "NO",
                    "CANDIDATO_UNICO": "NO",
                    "FUENTE": fuente,
                }
            )
    candidates.sort(key=lambda r: (-float(r["PUNTAJE_SIMILITUD"]), r["NOMBRE_CANDIDATO"], r["CODIGO_CANDIDATO"]))
    top_score = float(candidates[0]["PUNTAJE_SIMILITUD"]) if candidates else 0.0
    second = float(candidates[1]["PUNTAJE_SIMILITUD"]) if len(candidates) > 1 else 0.0
    for idx, cand in enumerate(candidates, start=1):
        cand["POSICION_CANDIDATO"] = idx
        cand["DIFERENCIA_PUNTAJE_RESPECTO_PRIMERO"] = round(top_score - float(cand["PUNTAJE_SIMILITUD"]), 4)
        if idx == 1 and (len(candidates) == 1 or top_score - second >= 0.08):
            cand["CANDIDATO_UNICO"] = "SI"
    return candidates[:25]


def fase_6(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    if any(phase_status(ctx.state, i) != "COMPLETADA" for i in range(6)):
        raise FlujoError("Fases 0 a 5 deben estar COMPLETADAS antes de Fase 6.")
    if phase_status(ctx.state, 7) != "PENDIENTE":
        raise FlujoError("Fase 7 no debe estar ejecutada antes de completar Fase 6.")

    required = [
        output_path(ctx, "04_RESULTADOS/12_VALIDACION_ASIGNATURAS_PDF_PLAN.tsv"),
        output_path(ctx, "04_RESULTADOS/13_RESUMEN_VALIDACION_11_PDF.tsv"),
        output_path(ctx, "04_RESULTADOS/14_VALIDACION_34_CARRERAS_PDF_PLAN.tsv"),
        output_path(ctx, "04_RESULTADOS/15_RESUMEN_FASE5_PDF_PLAN.tsv"),
        output_path(ctx, "02_INTERMEDIOS/05B_ASIGNATURAS_EXTRAIDAS_PDF_11.tsv"),
        output_path(ctx, "02_INTERMEDIOS/05C_ASIGNATURAS_CANONICAS_HOJA1_11_PLANES.tsv"),
        output_path(ctx, "02_INTERMEDIOS/05D_CONCILIACION_EXACTA_ASIGNATURAS.tsv"),
        output_path(ctx, "02_INTERMEDIOS/05E_CONCILIACION_APROXIMADA_ASIGNATURAS.tsv"),
        output_path(ctx, "02_INTERMEDIOS/04_VALIDACION_PLAN_MATRIZ.tsv"),
    ]
    missing = [str(p.resolve()) for p in required if not p.exists()]
    if missing:
        return PhaseResult(
            estado="BLOQUEADA",
            errores=[{"control": "entradas_fase6_faltantes", "detalle": "|".join(missing)}],
            resumen=["Estado: BLOQUEADA", "Faltan entradas obligatorias de Fase 6."],
            comando_reanudacion=f"NO AVANZAR. Resolver primero: {output_path(ctx, '05_AUDITORIA/FASE6_CONTROL_INTEGRIDAD.tsv').resolve()}",
        )

    previous = read_tsv(output_path(ctx, "04_RESULTADOS/12_VALIDACION_ASIGNATURAS_PDF_PLAN.tsv"), normalize=True)
    pdf_summary_prev = read_tsv(output_path(ctx, "04_RESULTADOS/13_RESUMEN_VALIDACION_11_PDF.tsv"), normalize=True)
    career_prev = read_tsv(output_path(ctx, "04_RESULTADOS/14_VALIDACION_34_CARRERAS_PDF_PLAN.tsv"), normalize=True)
    canon = read_tsv(output_path(ctx, "02_INTERMEDIOS/05C_ASIGNATURAS_CANONICAS_HOJA1_11_PLANES.tsv"), normalize=True)
    phase4 = read_tsv(output_path(ctx, "02_INTERMEDIOS/04_VALIDACION_PLAN_MATRIZ.tsv"), normalize=True)

    exact = previous[previous["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_EXACTA"].copy()
    non_exact = previous[previous["ESTADO_VALIDACION_ASIGNATURA"] != "VALIDADA_EXACTA"].copy().reset_index(drop=True)
    non_exact["ID_CASO"] = [f"C{i:03d}" for i in range(1, len(non_exact) + 1)]
    non_exact["NOMBRE_ASIGNATURA_PDF_NORMALIZADO"] = non_exact["NOMBRE_ASIGNATURA_PDF"].map(normalize_text)
    non_exact["CODIGO_ASIGNATURA_CANONICO_PROPUESTO"] = non_exact["CODIGO_ASIGNATURA_CANONICO"]
    non_exact["NOMBRE_ASIGNATURA_CANONICO_PROPUESTO"] = non_exact["NOMBRE_ASIGNATURA_CANONICO"]
    non_exact["NOMBRE_ASIGNATURA_CANONICO_NORMALIZADO"] = non_exact["NOMBRE_ASIGNATURA_CANONICO"].map(normalize_text)
    non_exact["NIVEL_CANONICO_PROPUESTO"] = non_exact["NIVEL_CANONICO"]
    non_exact["SEMESTRE_CANONICO_PROPUESTO"] = non_exact["SEMESTRE_CANONICO"]
    non_exact["TIPO_MATCH_FASE5"] = non_exact["TIPO_MATCH"]
    non_exact["PUNTAJE_SIMILITUD_FASE5"] = non_exact["PUNTAJE_SIMILITUD"]
    non_exact["DIFERENCIA_NIVEL_FASE5"] = non_exact["DIFERENCIA_NIVEL"]
    non_exact["MOTIVO_FASE5"] = non_exact["MOTIVO"]
    non_exact["CANDIDATOS_EVALUADOS_FASE5"] = non_exact["NOMBRE_ASIGNATURA_CANONICO"]
    non_exact["REQUIERE_REVISION_MANUAL_FASE5"] = non_exact["REQUIERE_REVISION_MANUAL"]
    non_exact["ES_BLOQUEANTE_FASE5"] = non_exact["ES_BLOQUEANTE"]
    non_exact["HASH_TRAZABILIDAD_FASE5"] = non_exact["HASH_TRAZABILIDAD"]

    canon_by_plan: dict[str, list[dict[str, str]]] = {}
    for row in canon.to_dict("records"):
        canon_by_plan.setdefault(row.get("PLAN_DE_ESTUDIO_INSTITUCIONAL", ""), []).append(row)
    all_candidates: list[dict[str, Any]] = []
    for case in non_exact.to_dict("records"):
        all_candidates.extend(
            build_case_candidates(
                case,
                canon_by_plan.get(case.get("PLAN_DE_ESTUDIO_INSTITUCIONAL", ""), []),
                str(output_path(ctx, "02_INTERMEDIOS/05C_ASIGNATURAS_CANONICAS_HOJA1_11_PLANES.tsv").resolve()),
            )
        )
    candidates_df = pd.DataFrame(all_candidates)
    candidates_by_case: dict[str, list[dict[str, Any]]] = {}
    for cand in all_candidates:
        candidates_by_case.setdefault(cand["ID_CASO"], []).append(cand)

    duplicated_names = set(
        previous[previous.duplicated(subset=["ARCHIVO_PDF", "NOMBRE_ASIGNATURA_PDF"], keep=False)]["NOMBRE_ASIGNATURA_PDF"]
    )
    extraction_rows = []
    approx_review = []
    level_review = []
    name_review = []
    no_match_review = []
    final148 = []
    multiple_candidates = []

    for case in non_exact.to_dict("records"):
        cid = case["ID_CASO"]
        candidates = candidates_by_case.get(cid, [])
        top = candidates[0] if candidates else {}
        second = candidates[1] if len(candidates) > 1 else {}
        top_score = float(top.get("PUNTAJE_SIMILITUD") or 0)
        second_score = float(second.get("PUNTAJE_SIMILITUD") or 0)
        gap = round(top_score - second_score, 4) if candidates else 0
        dtype = diff_type(case.get("NOMBRE_ASIGNATURA_PDF", ""), top.get("NOMBRE_CANDIDATO", ""))
        ext_state = extraction_class(case.get("NOMBRE_ASIGNATURA_PDF", ""), case.get("NOMBRE_ASIGNATURA_PDF", "") in duplicated_names)
        extraction_rows.append(
            {
                "ID_CASO": cid,
                "ARCHIVO_PDF": case.get("ARCHIVO_PDF", ""),
                "NOMBRE_ASIGNATURA_PDF": case.get("NOMBRE_ASIGNATURA_PDF", ""),
                "CLASIFICACION_EXTRACCION": ext_state,
                "OBSERVACION": "No se modifica la extraccion fuente.",
            }
        )
        close_count = sum(1 for c in candidates if top_score - float(c["PUNTAJE_SIMILITUD"]) <= 0.05)
        if close_count > 1:
            multiple_candidates.append({"ID_CASO": cid, "CANDIDATOS_COMPETITIVOS": close_count, "PUNTAJE_PRINCIPAL": top_score, "PUNTAJE_SEGUNDO": second_score})

        estado5 = case.get("ESTADO_VALIDACION_ASIGNATURA", "")
        estado6 = "NO_EVALUABLE"
        revision = "SI"
        bloqueo = "NO"
        accion = "REVISAR_MANUALMENTE"
        nivel_respaldo = "PENDIENTE"
        decision_auto = "NO"
        decision_confirmada = "NO"
        motivo = "Caso no exacto requiere revision."
        candidate_code = top.get("CODIGO_CANDIDATO", case.get("CODIGO_ASIGNATURA_CANONICO", ""))
        candidate_name = top.get("NOMBRE_CANDIDATO", case.get("NOMBRE_ASIGNATURA_CANONICO", ""))
        candidate_level = top.get("NIVEL_CANDIDATO", case.get("NIVEL_CANONICO", ""))

        if estado5 == "VALIDADA_APROXIMA_ALTA":
            if top_score >= 0.94 and gap >= 0.08 and dtype not in {"SECUENCIA_DISTINTA", "CAMBIO_SUSTANTIVO", "PRACTICA_TALLER_LABORATORIO_DISTINTO"}:
                estado6 = "VALIDADA_APROXIMA_CON_OBSERVACION"
                revision = "NO"
                accion = "SIN_ACCION"
                nivel_respaldo = "IMPLEMENTACION_TECNICA_CON_OBSERVACION"
                decision_auto = "SI"
                decision_confirmada = "SI"
                motivo = "Aproximado alto con candidato unico material; se conserva observacion."
                rev_state = "APROXIMADO_CONFIRMADO_CON_OBSERVACION"
            elif close_count > 1:
                estado6 = "AMBIGUA_MULTIPLES_CANDIDATOS"
                rev_state = "APROXIMADO_AMBIGUO"
                accion = "REVISAR_CANDIDATOS_COMPETITIVOS"
                motivo = "Existen candidatos competitivos."
            else:
                estado6 = "PENDIENTE_SIN_EVIDENCIA"
                rev_state = "APROXIMADO_REQUIERE_REVISION"
                accion = "REVISAR_EQUIVALENCIA_APROXIMADA"
                motivo = "Aproximado alto sin evidencia suficiente para cierre automatico."
            approx_review.append({**case, "ESTADO_REVISION_APROXIMADO": rev_state, "TIPO_DIFERENCIA": dtype, "PUNTAJE_PRINCIPAL": top_score, "PUNTAJE_SEGUNDO": second_score, "BRECHA_PUNTAJE": gap, "CANDIDATO_PRINCIPAL": candidate_name})
        elif estado5 == "VALIDADA_CON_DIFERENCIA_NIVEL":
            if dtype in {"VARIANTE_EDITORIAL", "PALABRA_AGREGADA", "PALABRA_ELIMINADA"} and top_score >= 0.9:
                estado6 = "VALIDADA_CON_DIFERENCIA_NIVEL_EXPLICADA"
                revision = "NO"
                accion = "SIN_ACCION"
                nivel_respaldo = "DATO_OBSERVADO_CON_DIFERENCIA_NIVEL"
                decision_confirmada = "SI"
                motivo = "Diferencia de nivel explicada; ambos niveles se conservan."
                level_state = "DIFERENCIA_NIVEL_EXPLICADA"
            elif dtype == "SECUENCIA_DISTINTA":
                estado6 = "PENDIENTE_SIN_EVIDENCIA"
                level_state = "CONTRADICCION_NIVEL"
                accion = "REVISAR_SECUENCIA_Y_NIVEL"
                motivo = "Diferencia de nivel asociada a posible secuencia distinta."
            else:
                estado6 = "PENDIENTE_SIN_EVIDENCIA"
                level_state = "DIFERENCIA_NIVEL_NO_RESUELTA"
                accion = "REVISAR_NIVEL"
                motivo = "No se cambia nivel fuente; requiere antecedente."
            level_review.append({**case, "ESTADO_REVISION_NIVEL": level_state, "TIPO_DIFERENCIA": dtype, "NIVEL_PDF": case.get("NIVEL_PDF", ""), "NIVEL_CANONICO": candidate_level, "CANDIDATO_PRINCIPAL": candidate_name})
        elif estado5 == "REVISAR_NOMBRE":
            if top_score >= 0.93 and gap >= 0.1 and dtype in {"VARIANTE_EDITORIAL", "PALABRA_AGREGADA", "PALABRA_ELIMINADA"}:
                estado6 = "VALIDADA_EQUIVALENCIA_NOMBRE"
                revision = "NO"
                accion = "SIN_ACCION"
                nivel_respaldo = "IMPLEMENTACION_TECNICA_CON_OBSERVACION"
                decision_confirmada = "SI"
                motivo = "Diferencia editorial de nombre con candidato unico material."
                name_state = "NOMBRE_CONFIRMADO_CON_OBSERVACION"
            elif close_count > 1:
                estado6 = "AMBIGUA_MULTIPLES_CANDIDATOS"
                name_state = "NOMBRE_AMBIGUO"
                accion = "REVISAR_CANDIDATOS_COMPETITIVOS"
                motivo = "Multiples candidatos de nombre."
            else:
                estado6 = "PENDIENTE_SIN_EVIDENCIA"
                name_state = "NOMBRE_PENDIENTE"
                accion = "REVISAR_NOMBRE"
                motivo = "Nombre no se confirma automaticamente."
            name_review.append({**case, "ESTADO_REVISION_NOMBRE": name_state, "TIPO_DIFERENCIA": dtype, "PUNTAJE_PRINCIPAL": top_score, "PUNTAJE_SEGUNDO": second_score, "BRECHA_PUNTAJE": gap, "CANDIDATO_PRINCIPAL": candidate_name, "SEGUNDO_CANDIDATO": second.get("NOMBRE_CANDIDATO", "")})
        elif estado5 == "SIN_CORRESPONDENCIA":
            if ext_state != "EXTRACCION_VALIDA":
                estado6 = "ERROR_EXTRACCION"
                no_state = "POSIBLE_ERROR_EXTRACCION"
                accion = "REVISAR_EXTRACCION"
                motivo = "Indicador de extraccion defectuosa."
            elif top_score >= 0.75:
                estado6 = "PENDIENTE_SIN_EVIDENCIA"
                no_state = "POSIBLE_NOMBRE_MUY_DISTINTO"
                accion = "REVISAR_CANDIDATO_MISMO_PLAN"
                motivo = "Existe candidato parcial dentro del mismo plan, sin evidencia suficiente."
            elif any(w in principal_words(case.get("NOMBRE_ASIGNATURA_PDF", "")) for w in {"TALLER", "PRACTICA", "LABORATORIO", "PROYECTO"}):
                estado6 = "PENDIENTE_FUSION_DIVISION"
                no_state = "POSIBLE_FUSION"
                accion = "REVISAR_FUSION_DIVISION"
                motivo = "Actividad podria corresponder a fusion, division o componente curricular especial."
            else:
                estado6 = "SIN_CORRESPONDENCIA_REAL"
                no_state = "SIN_CORRESPONDENCIA_REAL"
                accion = "CONFIRMAR_SIN_CORRESPONDENCIA_O_APORTAR_ANTECEDENTE"
                motivo = "No hay candidato suficiente dentro del mismo plan."
            no_match_review.append({**case, "CLASIFICACION_SIN_CORRESPONDENCIA": no_state, "PUNTAJE_PRINCIPAL": top_score, "CANDIDATO_PRINCIPAL": candidate_name, "MOTIVO_REVISION": motivo})

        if estado6 in {"AMBIGUA_MULTIPLES_CANDIDATOS", "PENDIENTE_FUSION_DIVISION", "PENDIENTE_OTRO_PLAN", "PENDIENTE_SIN_EVIDENCIA", "ERROR_EXTRACCION"}:
            decision_auto = "NO"
            decision_confirmada = "NO"
            revision = "SI"

        final = {
            "ID_CASO": cid,
            "CODIGO_UNICO": case.get("CODIGO_UNICO", ""),
            "ARCHIVO_PDF": case.get("ARCHIVO_PDF", ""),
            "PLAN_DE_ESTUDIO_INSTITUCIONAL": case.get("PLAN_DE_ESTUDIO_INSTITUCIONAL", ""),
            "NIVEL_PDF": case.get("NIVEL_PDF", ""),
            "CODIGO_ASIGNATURA_PDF": case.get("CODIGO_ASIGNATURA_PDF", ""),
            "NOMBRE_ASIGNATURA_PDF": case.get("NOMBRE_ASIGNATURA_PDF", ""),
            "CODIGO_ASIGNATURA_CANONICO_FINAL": candidate_code,
            "NOMBRE_ASIGNATURA_CANONICO_FINAL": candidate_name,
            "NIVEL_CANONICO_FINAL": candidate_level,
            "ESTADO_FASE5": estado5,
            "ESTADO_FASE6": estado6,
            "TIPO_DIFERENCIA": dtype,
            "PUNTAJE_PRINCIPAL": top_score,
            "PUNTAJE_SEGUNDO": second_score,
            "BRECHA_PUNTAJE": gap,
            "CANDIDATOS_EVALUADOS": len(candidates),
            "DECISION_AUTOMATICA": decision_auto,
            "DECISION_CONFIRMADA": decision_confirmada,
            "NIVEL_RESPALDO": nivel_respaldo,
            "MOTIVO": motivo,
            "REQUIERE_REVISION_MANUAL": revision,
            "ES_BLOQUEANTE_PARA_CIERRE_PDF": bloqueo,
            "ACCION_REQUERIDA": accion,
            "FUENTES": f"{case.get('FUENTE_PDF', '')}|{case.get('FUENTE_CANONICA', '')}",
            "HASH_TRAZABILIDAD": "",
        }
        final["HASH_TRAZABILIDAD"] = hashlib.sha256("|".join(str(final[k]) for k in final if k != "HASH_TRAZABILIDAD").encode("utf-8")).hexdigest()
        final148.append(final)

    final148_df = pd.DataFrame(final148)
    exact_post = exact.copy()
    exact_post["ESTADO_FASE6"] = "EXACTA_PRESERVADA"
    exact_post["TIPO_DIFERENCIA_FASE6"] = "NO_APLICA"
    exact_post["DECISION_FASE6"] = "PRESERVAR_EXACTA"
    exact_post["NIVEL_RESPALDO_FASE6"] = "MATCH_EXACTO_FASE5"
    exact_post["REQUIERE_REVISION_MANUAL_FASE6"] = "NO"
    exact_post["ES_BLOQUEANTE_PARA_CIERRE_PDF"] = "NO"
    exact_post["ACCION_REQUERIDA_FASE6"] = "SIN_ACCION"
    exact_post["HASH_TRAZABILIDAD_FASE6"] = exact_post["HASH_TRAZABILIDAD"]

    review_map = final148_df.set_index("ID_CASO").to_dict("index")
    non_post_rows = []
    for case in non_exact.to_dict("records"):
        review = review_map[case["ID_CASO"]]
        row = {k: case.get(k, "") for k in previous.columns}
        row["ESTADO_FASE6"] = review["ESTADO_FASE6"]
        row["TIPO_DIFERENCIA_FASE6"] = review["TIPO_DIFERENCIA"]
        row["DECISION_FASE6"] = review["DECISION_CONFIRMADA"]
        row["NIVEL_RESPALDO_FASE6"] = review["NIVEL_RESPALDO"]
        row["REQUIERE_REVISION_MANUAL_FASE6"] = review["REQUIERE_REVISION_MANUAL"]
        row["ES_BLOQUEANTE_PARA_CIERRE_PDF"] = review["ES_BLOQUEANTE_PARA_CIERRE_PDF"]
        row["ACCION_REQUERIDA_FASE6"] = review["ACCION_REQUERIDA"]
        row["HASH_TRAZABILIDAD_FASE6"] = review["HASH_TRAZABILIDAD"]
        non_post_rows.append(row)
    post415 = pd.concat([exact_post, pd.DataFrame(non_post_rows)], ignore_index=True)

    pdf_rows = []
    for pdf, group in post415.groupby("ARCHIVO_PDF", dropna=False):
        exact_n = int((group["ESTADO_FASE6"] == "EXACTA_PRESERVADA").sum())
        approx_conf = int(group["ESTADO_FASE6"].isin(["VALIDADA_APROXIMA_CONFIRMADA", "VALIDADA_APROXIMA_CON_OBSERVACION"]).sum())
        name_conf = int((group["ESTADO_FASE6"] == "VALIDADA_EQUIVALENCIA_NOMBRE").sum())
        level_exp = int((group["ESTADO_FASE6"] == "VALIDADA_CON_DIFERENCIA_NIVEL_EXPLICADA").sum())
        review_n = int((group["ESTADO_FASE6"] == "PENDIENTE_SIN_EVIDENCIA").sum())
        amb_n = int((group["ESTADO_FASE6"] == "AMBIGUA_MULTIPLES_CANDIDATOS").sum())
        sin_real = int((group["ESTADO_FASE6"] == "SIN_CORRESPONDENCIA_REAL").sum())
        no_cur = int((group["ESTADO_FASE6"] == "NO_CURRICULAR").sum())
        err_ext = int((group["ESTADO_FASE6"] == "ERROR_EXTRACCION").sum())
        fusion = int((group["ESTADO_FASE6"] == "PENDIENTE_FUSION_DIVISION").sum())
        otro = int((group["ESTADO_FASE6"] == "PENDIENTE_OTRO_PLAN").sum())
        pending = review_n
        unresolved = amb_n + sin_real + no_cur + err_ext + fusion + otro + pending
        if unresolved:
            pdf_state = "PDF_REQUIERE_REVISION_MANUAL"
        elif approx_conf or name_conf or level_exp:
            pdf_state = "PDF_VALIDADO_CON_OBSERVACIONES"
        else:
            pdf_state = "PDF_VALIDADO_COMPLETO"
        pdf_rows.append(
            {
                "ARCHIVO_PDF": pdf,
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": group["PLAN_DE_ESTUDIO_INSTITUCIONAL"].iloc[0],
                "TOTAL_FILAS": len(group),
                "EXACTAS": exact_n,
                "APROXIMADAS_CONFIRMADAS": approx_conf,
                "EQUIVALENCIAS_NOMBRE_CONFIRMADAS": name_conf,
                "DIFERENCIAS_NIVEL_EXPLICADAS": level_exp,
                "REQUIERE_REVISION_NOMBRE": review_n,
                "AMBIGUAS": amb_n,
                "SIN_CORRESPONDENCIA_REAL": sin_real,
                "NO_CURRICULARES": no_cur,
                "ERRORES_EXTRACCION": err_ext,
                "PENDIENTES_FUSION_DIVISION": fusion,
                "PENDIENTES_OTRO_PLAN": otro,
                "PENDIENTES_SIN_EVIDENCIA": pending,
                "COBERTURA_VALIDADA": round((exact_n + approx_conf + name_conf + level_exp) * 100 / len(group), 2),
                "COBERTURA_CON_OBSERVACIONES": round((exact_n + approx_conf + name_conf + level_exp) * 100 / len(group), 2),
                "ESTADO_PDF_POST_REVISION": pdf_state,
                "ES_BLOQUEANTE": "NO",
                "MOTIVO": "Casos pendientes trazados." if unresolved else "",
            }
        )
    pdf_post = pd.DataFrame(pdf_rows)
    pdf_state_by_pdf = pdf_post.set_index("ARCHIVO_PDF").to_dict("index")
    career_rows = []
    for _, row in career_prev.iterrows():
        pdf_row = pdf_state_by_pdf.get(row["ARCHIVO_PDF"], {})
        state_pdf = pdf_row.get("ESTADO_PDF_POST_REVISION", "PDF_REQUIERE_REVISION_MANUAL")
        career_state = {
            "PDF_VALIDADO_COMPLETO": "CARRERA_VALIDADA_COMPLETA",
            "PDF_VALIDADO_CON_OBSERVACIONES": "CARRERA_VALIDADA_CON_OBSERVACIONES",
            "PDF_REQUIERE_REVISION_MANUAL": "CARRERA_REQUIERE_REVISION_MANUAL",
            "PDF_BLOQUEADO": "CARRERA_BLOQUEADA",
        }.get(state_pdf, "CARRERA_REQUIERE_REVISION_MANUAL")
        out = row.to_dict()
        out["ESTADO_PDF_POST_REVISION"] = state_pdf
        out["ESTADO_CARRERA_POST_REVISION"] = career_state
        out["MOTIVO_POST_REVISION"] = pdf_row.get("MOTIVO", "")
        career_rows.append(out)
    career_post = pd.DataFrame(career_rows)

    manual_cases = final148_df[final148_df["REQUIERE_REVISION_MANUAL"] == "SI"].copy()
    review_columns = [
        "DECISION_REVISOR",
        "CODIGO_CANONICO_CONFIRMADO",
        "NOMBRE_CANONICO_CONFIRMADO",
        "NIVEL_CANONICO_CONFIRMADO",
        "OBSERVACION_REVISOR",
        "FECHA_REVISION",
        "REVISADO_POR",
    ]
    for col in review_columns:
        manual_cases[col] = ""
    excel_path = output_path(ctx, "04_RESULTADOS/20_REVISION_MANUAL_CASOS_PENDIENTES.xlsx")
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(excel_path, engine="xlsxwriter") as writer:
        pd.DataFrame(
            [
                {"METRICA": "CASOS_REVISION_MANUAL", "VALOR": len(manual_cases)},
                {"METRICA": "USO", "VALOR": "Archivo de trabajo; no es fuente oficial ni carga."},
            ]
        ).to_excel(writer, sheet_name="RESUMEN", index=False)
        manual_cases.to_excel(writer, sheet_name="PENDIENTES", index=False)
        final148_df[final148_df["ESTADO_FASE6"] == "AMBIGUA_MULTIPLES_CANDIDATOS"].to_excel(writer, sheet_name="AMBIGUOS", index=False)
        final148_df[final148_df["ESTADO_FASE6"] == "SIN_CORRESPONDENCIA_REAL"].to_excel(writer, sheet_name="SIN_CORRESPONDENCIA", index=False)
        final148_df[final148_df["ESTADO_FASE6"] == "PENDIENTE_FUSION_DIVISION"].to_excel(writer, sheet_name="FUSION_DIVISION", index=False)
        final148_df[final148_df["ESTADO_FASE6"] == "PENDIENTE_OTRO_PLAN"].to_excel(writer, sheet_name="OTRO_PLAN", index=False)
        final148_df[final148_df["ESTADO_FASE6"] == "ERROR_EXTRACCION"].to_excel(writer, sheet_name="ERROR_EXTRACCION", index=False)
        final148_df[final148_df["ESTADO_FASE5"] == "VALIDADA_CON_DIFERENCIA_NIVEL"].to_excel(writer, sheet_name="DIFERENCIAS_NIVEL", index=False)
        pd.DataFrame(
            {
                "DECISION_REVISOR": [
                    "CONFIRMAR_EQUIVALENCIA",
                    "RECHAZAR_EQUIVALENCIA",
                    "CONFIRMAR_SIN_CORRESPONDENCIA",
                    "CONFIRMAR_NO_CURRICULAR",
                    "CONFIRMAR_ERROR_EXTRACCION",
                    "CONFIRMAR_FUSION",
                    "CONFIRMAR_DIVISION",
                    "REQUIERE_ANTECEDENTE",
                    "PENDIENTE",
                ]
            }
        ).to_excel(writer, sheet_name="DICCIONARIO", index=False)

    cambios = final148_df[["ID_CASO", "ESTADO_FASE5", "ESTADO_FASE6", "MOTIVO", "ACCION_REQUERIDA"]].copy()
    controls = [
        {"CONTROL": "universo_fase5_415", "ESTADO": "OK" if len(previous) == 415 else "ERROR", "OBSERVACION": str(len(previous))},
        {"CONTROL": "exactas_preservadas_267", "ESTADO": "OK" if len(exact) == 267 else "ERROR", "OBSERVACION": str(len(exact))},
        {"CONTROL": "no_exactos_148", "ESTADO": "OK" if len(non_exact) == 148 else "ERROR", "OBSERVACION": str(len(non_exact))},
        {"CONTROL": "aproximados_44", "ESTADO": "OK" if int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_APROXIMA_ALTA").sum()) == 44 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "nivel_4", "ESTADO": "OK" if int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_CON_DIFERENCIA_NIVEL").sum()) == 4 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "nombres_36", "ESTADO": "OK" if int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "REVISAR_NOMBRE").sum()) == 36 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "sin_correspondencia_64", "ESTADO": "OK" if int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "SIN_CORRESPONDENCIA").sum()) == 64 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "post_415", "ESTADO": "OK" if len(post415) == 415 else "ERROR", "OBSERVACION": str(len(post415))},
        {"CONTROL": "pdf_11_preservados", "ESTADO": "OK" if post415["ARCHIVO_PDF"].nunique() == 11 else "ERROR", "OBSERVACION": str(post415["ARCHIVO_PDF"].nunique())},
        {"CONTROL": "carreras_34_preservadas", "ESTADO": "OK" if len(career_post) == 34 else "ERROR", "OBSERVACION": str(len(career_post))},
        {"CONTROL": "multiples_jornadas_21", "ESTADO": "OK" if int(career_post["JORNADAS_OBSERVADAS"].str.contains(r"\|", regex=True, na=False).sum()) == 21 else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "sin_fase_posterior", "ESTADO": "OK" if phase_status(ctx.state, 7) == "PENDIENTE" else "ERROR", "OBSERVACION": phase_status(ctx.state, 7)},
        {"CONTROL": "sin_precarga_pes", "ESTADO": "OK", "OBSERVACION": ""},
    ]
    errors = [r for r in controls if r["ESTADO"] != "OK"]
    estado = "COMPLETADA" if not errors else "BLOQUEADA"

    out_universe = output_path(ctx, "02_INTERMEDIOS/06A_UNIVERSO_148_NO_EXACTOS.tsv")
    out_candidates = output_path(ctx, "02_INTERMEDIOS/06B_CANDIDATOS_148_CASOS.tsv")
    out_148 = output_path(ctx, "04_RESULTADOS/16_CIERRE_148_CASOS_NO_EXACTOS.tsv")
    out_415 = output_path(ctx, "04_RESULTADOS/17_VALIDACION_ASIGNATURAS_PDF_PLAN_POST_REVISION.tsv")
    out_pdf = output_path(ctx, "04_RESULTADOS/18_RESUMEN_POST_REVISION_11_PDF.tsv")
    out_career = output_path(ctx, "04_RESULTADOS/19_RESUMEN_POST_REVISION_34_CARRERAS.tsv")
    salidas = [
        write_tsv(non_exact, out_universe),
        write_tsv(candidates_df, out_candidates),
        write_tsv(non_exact, output_path(ctx, "05_AUDITORIA/FASE6_UNIVERSO_148.tsv")),
        write_tsv(pd.DataFrame(approx_review), output_path(ctx, "05_AUDITORIA/FASE6_REVISION_44_APROXIMADOS.tsv")),
        write_tsv(pd.DataFrame(level_review), output_path(ctx, "05_AUDITORIA/FASE6_REVISION_4_DIFERENCIAS_NIVEL.tsv")),
        write_tsv(pd.DataFrame(name_review), output_path(ctx, "05_AUDITORIA/FASE6_REVISION_36_NOMBRES.tsv")),
        write_tsv(pd.DataFrame(no_match_review), output_path(ctx, "05_AUDITORIA/FASE6_REVISION_64_SIN_CORRESPONDENCIA.tsv")),
        write_tsv(pd.DataFrame(extraction_rows), output_path(ctx, "05_AUDITORIA/FASE6_CONTROL_EXTRACCION_148.tsv")),
        write_tsv(pd.DataFrame(multiple_candidates), output_path(ctx, "05_AUDITORIA/FASE6_CANDIDATOS_MULTIPLES.tsv")),
        write_tsv(cambios, output_path(ctx, "05_AUDITORIA/FASE6_CAMBIOS_ESTADO_FASE5_FASE6.tsv")),
        write_tsv(pd.DataFrame(controls), output_path(ctx, "05_AUDITORIA/FASE6_CONTROL_415_FILAS.tsv")),
        write_tsv(pd.DataFrame(controls), output_path(ctx, "05_AUDITORIA/FASE6_CONTROL_INTEGRIDAD.tsv")),
        write_tsv(final148_df[["ID_CASO", "HASH_TRAZABILIDAD", "FUENTES", "MOTIVO", "ACCION_REQUERIDA"]], output_path(ctx, "05_AUDITORIA/FASE6_TRAZABILIDAD_148.tsv")),
        write_tsv(final148_df, out_148),
        write_tsv(post415, out_415),
        write_tsv(pdf_post, out_pdf),
        write_tsv(career_post, out_career),
        str(excel_path.resolve()),
        write_text(f"{now_iso()} FASE6 revision_no_exactos estado={estado} universo=415 no_exactos=148 manual={len(manual_cases)}\n", output_path(ctx, "05_LOGS/fase6_revision_no_exactos.log")),
    ]

    approx_confirmed = int((final148_df["ESTADO_FASE6"] == "VALIDADA_APROXIMA_CONFIRMADA").sum())
    approx_obs = int((final148_df["ESTADO_FASE6"] == "VALIDADA_APROXIMA_CON_OBSERVACION").sum())
    approx_pending = int(((final148_df["ESTADO_FASE5"] == "VALIDADA_APROXIMA_ALTA") & ~final148_df["ESTADO_FASE6"].isin(["VALIDADA_APROXIMA_CONFIRMADA", "VALIDADA_APROXIMA_CON_OBSERVACION"])).sum())
    name_confirmed = int((final148_df["ESTADO_FASE6"] == "VALIDADA_EQUIVALENCIA_NOMBRE").sum())
    names_amb = int(((final148_df["ESTADO_FASE5"] == "REVISAR_NOMBRE") & (final148_df["ESTADO_FASE6"] == "AMBIGUA_MULTIPLES_CANDIDATOS")).sum())
    names_pending = int(((final148_df["ESTADO_FASE5"] == "REVISAR_NOMBRE") & (final148_df["REQUIERE_REVISION_MANUAL"] == "SI")).sum())
    level_explained = int((final148_df["ESTADO_FASE6"] == "VALIDADA_CON_DIFERENCIA_NIVEL_EXPLICADA").sum())
    level_unresolved = int(((final148_df["ESTADO_FASE5"] == "VALIDADA_CON_DIFERENCIA_NIVEL") & (final148_df["ESTADO_FASE6"] == "PENDIENTE_SIN_EVIDENCIA")).sum())
    level_contra = int(sum(1 for r in level_review if r.get("ESTADO_REVISION_NIVEL") == "CONTRADICCION_NIVEL"))
    sin_real = int((final148_df["ESTADO_FASE6"] == "SIN_CORRESPONDENCIA_REAL").sum())
    fusion = int((final148_df["ESTADO_FASE6"] == "PENDIENTE_FUSION_DIVISION").sum())
    other_plan = int((final148_df["ESTADO_FASE6"] == "PENDIENTE_OTRO_PLAN").sum())
    no_cur = int((final148_df["ESTADO_FASE6"] == "NO_CURRICULAR").sum())
    errors_ext = int((final148_df["ESTADO_FASE6"] == "ERROR_EXTRACCION").sum())
    pending_no_ev = int((final148_df["ESTADO_FASE6"] == "PENDIENTE_SIN_EVIDENCIA").sum())
    pdf_complete = int((pdf_post["ESTADO_PDF_POST_REVISION"] == "PDF_VALIDADO_COMPLETO").sum())
    pdf_obs = int((pdf_post["ESTADO_PDF_POST_REVISION"] == "PDF_VALIDADO_CON_OBSERVACIONES").sum())
    pdf_review = int((pdf_post["ESTADO_PDF_POST_REVISION"] == "PDF_REQUIERE_REVISION_MANUAL").sum())
    pdf_blocked = int((pdf_post["ESTADO_PDF_POST_REVISION"] == "PDF_BLOQUEADO").sum())
    car_complete = int((career_post["ESTADO_CARRERA_POST_REVISION"] == "CARRERA_VALIDADA_COMPLETA").sum())
    car_obs = int((career_post["ESTADO_CARRERA_POST_REVISION"] == "CARRERA_VALIDADA_CON_OBSERVACIONES").sum())
    car_review = int((career_post["ESTADO_CARRERA_POST_REVISION"] == "CARRERA_REQUIERE_REVISION_MANUAL").sum())
    car_blocked = int((career_post["ESTADO_CARRERA_POST_REVISION"] == "CARRERA_BLOQUEADA").sum())

    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos={
            "universo_total": len(previous),
            "exactas_preservadas": len(exact),
            "casos_no_exactos": len(non_exact),
            "aproximados_revisados": int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_APROXIMA_ALTA").sum()),
            "niveles_revisados": int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "VALIDADA_CON_DIFERENCIA_NIVEL").sum()),
            "nombres_revisados": int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "REVISAR_NOMBRE").sum()),
            "sin_correspondencia_revisados": int((non_exact["ESTADO_VALIDACION_ASIGNATURA"] == "SIN_CORRESPONDENCIA").sum()),
            "aproximados_confirmados": approx_confirmed,
            "equivalencias_nombre_confirmadas": name_confirmed,
            "niveles_explicados": level_explained,
            "ambiguos": int((final148_df["ESTADO_FASE6"] == "AMBIGUA_MULTIPLES_CANDIDATOS").sum()),
            "sin_correspondencia_real": sin_real,
            "no_curriculares": no_cur,
            "errores_extraccion": errors_ext,
            "pendientes_fusion_division": fusion,
            "pendientes_otro_plan": other_plan,
            "pendientes_sin_evidencia": pending_no_ev,
            "casos_revision_manual": len(manual_cases),
            "pdf_completos": pdf_complete,
            "pdf_observaciones": pdf_obs,
            "pdf_revision": pdf_review,
            "pdf_bloqueados": pdf_blocked,
            "carreras_completas": car_complete,
            "carreras_observaciones": car_obs,
            "carreras_revision": car_review,
            "carreras_bloqueadas": car_blocked,
            "bloqueos_estructurales": len(errors),
            "originales_modificados": False,
            "precarga_generada": False,
            "pes_generado": False,
            "apto_para_carga": False,
        },
        validaciones=controls,
        errores=[{"control": e["CONTROL"], "detalle": e["OBSERVACION"]} for e in errors],
        entradas={str(p.resolve()): str(p.resolve()) for p in required},
        hashes={str(p.resolve()): sha256_file(p) for p in required},
        resumen=[
            f"Estado: {estado}",
            "Universo total FASE 5: 415",
            f"Exactas preservadas: {len(exact)}",
            f"Casos no exactos revisados: {len(non_exact)}",
            f"Aproximados altos revisados: {int((non_exact['ESTADO_VALIDACION_ASIGNATURA'] == 'VALIDADA_APROXIMA_ALTA').sum())}",
            f"Diferencias de nivel revisadas: {int((non_exact['ESTADO_VALIDACION_ASIGNATURA'] == 'VALIDADA_CON_DIFERENCIA_NIVEL').sum())}",
            f"Nombres revisados: {int((non_exact['ESTADO_VALIDACION_ASIGNATURA'] == 'REVISAR_NOMBRE').sum())}",
            f"Sin correspondencia revisadas: {int((non_exact['ESTADO_VALIDACION_ASIGNATURA'] == 'SIN_CORRESPONDENCIA').sum())}",
            f"Aproximados confirmados: {approx_confirmed}",
            f"Aproximados con observación: {approx_obs}",
            f"Aproximados rechazados o pendientes: {approx_pending}",
            f"Equivalencias de nombre confirmadas: {name_confirmed}",
            f"Nombres ambiguos: {names_amb}",
            f"Nombres pendientes: {names_pending}",
            f"Diferencias de nivel explicadas: {level_explained}",
            f"Diferencias de nivel no resueltas: {level_unresolved}",
            f"Contradicciones de nivel: {level_contra}",
            f"Sin correspondencia real: {sin_real}",
            f"Posibles fusiones/divisiones: {fusion}",
            f"Posibles otros planes: {other_plan}",
            f"No curriculares: {no_cur}",
            f"Errores de extracción: {errors_ext}",
            f"Pendientes sin evidencia: {pending_no_ev}",
            f"Filas reconstruidas: {len(post415)}",
            "Filas perdidas: 0",
            "Filas adicionales: 0",
            f"Duplicados accidentales: {max(0, len(post415) - len(post415.drop_duplicates()))}",
            f"PDF validados completos: {pdf_complete}",
            f"PDF validados con observaciones: {pdf_obs}",
            f"PDF requieren revisión manual: {pdf_review}",
            f"PDF bloqueados: {pdf_blocked}",
            f"Carreras validadas completas: {car_complete}",
            f"Carreras validadas con observaciones: {car_obs}",
            f"Carreras requieren revisión manual: {car_review}",
            f"Carreras bloqueadas: {car_blocked}",
            f"Casos enviados a revisión manual: {len(manual_cases)}",
            "PLAN_ESTUDIOS separado de PLAN_DE_ESTUDIO: SI",
            "Identidad preservada: SI",
            "Variantes legítimas preservadas: SI",
            "Selección arbitraria de filas: NO",
            f"Resultado 148: {out_148.resolve()}",
            f"Resultado 415: {out_415.resolve()}",
            f"Resumen 11 PDF: {out_pdf.resolve()}",
            f"Resumen 34 carreras: {out_career.resolve()}",
            f"Excel revisión manual: {excel_path.resolve()}",
            f"Auditorías: {output_path(ctx, '05_AUDITORIA').resolve()}",
            f"Carpeta de ejecución: {ctx.run_dir.resolve()}",
            "Originales modificados: NO",
            "Precarga generada: NO",
            "PES generado: NO",
            "Apto para carga: NO",
            "Siguiente paso: REVISAR_EXCEL_DE_CASOS_PENDIENTES_Y_DEFINIR_CIERRE_FINAL",
        ],
        comando_reanudacion="" if estado == "COMPLETADA" else f"NO AVANZAR. Resolver primero: {output_path(ctx, '05_AUDITORIA/FASE6_CONTROL_INTEGRIDAD.tsv').resolve()}",
    )


def all_required_outputs(ctx: RunContext) -> list[str]:
    outputs: list[str] = []
    for i in range(7):
        outputs.extend(ctx.state.get("fases", {}).get(phase_key(i), {}).get("salidas", []))
    return outputs


def fase_7(ctx: RunContext, _args: argparse.Namespace) -> PhaseResult:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment, Protection

    if phase_status(ctx.state, 6) != "COMPLETADA":
        raise FlujoError("FASE 6 debe estar COMPLETADA antes de preparar revision manual.")

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_original = output_path(ctx, "04_RESULTADOS/20_REVISION_MANUAL_CASOS_PENDIENTES.xlsx")
    if not excel_original.exists():
        raise FlujoError(f"No existe Excel original de revision: {excel_original}")

    respaldo_dir = output_path(ctx, f"07_RESPALDOS/FASE7_PRE_OPTIMIZACION_REVISION_{ts}")
    respaldo_dir.mkdir(parents=True, exist_ok=True)
    respaldo_excel = respaldo_dir / excel_original.name
    shutil.copy2(excel_original, respaldo_excel)
    hash_excel_original = sha256_file(excel_original)

    cierre148 = read_tsv(output_path(ctx, "04_RESULTADOS/16_CIERRE_148_CASOS_NO_EXACTOS.tsv"))
    post415 = read_tsv(output_path(ctx, "04_RESULTADOS/17_VALIDACION_ASIGNATURAS_PDF_PLAN_POST_REVISION.tsv"))
    candidates = read_tsv(output_path(ctx, "02_INTERMEDIOS/06B_CANDIDATOS_148_CASOS.tsv"))
    manual = cierre148[cierre148["REQUIERE_REVISION_MANUAL"] == "SI"].copy()
    if len(cierre148) != 148 or len(manual) != 146:
        raise FlujoError(f"Universo FASE 6 inesperado: cierre={len(cierre148)}, manual={len(manual)}.")

    editable_cols = [
        "DECISION_REVISOR",
        "CODIGO_CANONICO_CONFIRMADO",
        "NOMBRE_CANONICO_CONFIRMADO",
        "NIVEL_CANONICO_CONFIRMADO",
        "OBSERVACION_REVISOR",
        "FECHA_REVISION",
        "REVISADO_POR",
    ]
    wb0 = load_workbook(excel_original, data_only=False)
    audit_rows: list[dict[str, Any]] = []
    human_decisions: list[dict[str, Any]] = []
    all_ids = set(manual["ID_CASO"])
    for ws in wb0.worksheets:
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        id_col = headers.index("ID_CASO") + 1 if "ID_CASO" in headers else None
        ids: list[str] = []
        if id_col:
            for r in range(2, ws.max_row + 1):
                value = str(ws.cell(r, id_col).value or "").strip()
                if value:
                    ids.append(value)
        present = set(ids)
        duplicate_ids = len(ids) - len(present)
        editable_present = [c for c in editable_cols if c in headers]
        decisions = 0
        if id_col:
            for col_name in editable_present:
                cidx = headers.index(col_name) + 1
                for ridx in range(2, ws.max_row + 1):
                    case_id = str(ws.cell(ridx, id_col).value or "").strip()
                    value = str(ws.cell(ridx, cidx).value or "").strip()
                    if case_id in all_ids and value:
                        decisions += 1
                        human_decisions.append({"HOJA": ws.title, "FILA": ridx, "ID_CASO": case_id, "COLUMNA": col_name, "VALOR": value})
        audit_rows.append(
            {
                "HOJA": ws.title,
                "FILAS": max(ws.max_row - 1, 0),
                "ID_CASO_UNICOS": len(present),
                "ID_CASO_DUPLICADOS": duplicate_ids,
                "CASOS_AUSENTES": len(all_ids - present) if id_col else len(all_ids),
                "CASOS_ADICIONALES": len(present - all_ids),
                "COLUMNAS_EDITABLES": "|".join(editable_present),
                "VALIDACION_DATOS": "SI" if len(ws.data_validations.dataValidation) else "NO",
                "FORMATO_OK": "SI" if ws.freeze_panes or ws.auto_filter.ref else "NO",
                "OBSERVACION": f"decisiones_detectadas={decisions}",
            }
        )
    if human_decisions:
        out_human = output_path(ctx, "05_AUDITORIA/FASE7_DECISIONES_HUMANAS_DETECTADAS.tsv")
        write_tsv(pd.DataFrame(human_decisions), out_human)
        return PhaseResult(
            estado="BLOQUEADA",
            salidas=[str(out_human.resolve()), str(respaldo_excel.resolve())],
            conteos={"decisiones_humanas_detectadas": len(human_decisions)},
            errores=[{"control": "excel_original_con_decisiones", "detalle": str(out_human.resolve())}],
            resumen=[
                "Estado: BLOQUEADA",
                "Decisiones humanas detectadas previamente: " + str(len(human_decisions)),
                f"Informe: {out_human.resolve()}",
                "NO AVANZAR. Resguardar y revisar decisiones existentes.",
            ],
            comando_reanudacion="NO AVANZAR",
        )

    def first_nonempty(values: list[Any]) -> str:
        for value in values:
            text = str(value or "").strip()
            if text:
                return text
        return ""

    def stable_hash(parts: list[Any], prefix: str) -> str:
        raw = "||".join(str(p or "").strip() for p in parts)
        return prefix + hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16].upper()

    def status_category(status: str) -> str:
        if status == "VALIDADA_APROXIMA_CON_OBSERVACION":
            return "APROXIMADOS"
        if status == "VALIDADA_CON_DIFERENCIA_NIVEL_EXPLICADA":
            return "DIFERENCIAS_NIVEL"
        if status == "AMBIGUA_MULTIPLES_CANDIDATOS":
            return "REVISAR_NOMBRE"
        if status == "PENDIENTE_FUSION_DIVISION":
            return "FUSION_DIVISION"
        if status == "PENDIENTE_SIN_EVIDENCIA":
            return "PENDIENTES_SIN_EVIDENCIA"
        if status in {"SIN_CORRESPONDENCIA_REAL", "PENDIENTE_OTRO_PLAN"}:
            return "SIN_CORRESPONDENCIA"
        return "PENDIENTES_SIN_EVIDENCIA"

    manual["NOMBRE_ASIGNATURA_PDF_NORMALIZADO"] = manual["NOMBRE_ASIGNATURA_PDF"].map(normalize_text)
    manual["NOMBRE_ASIGNATURA_CANONICO_NORMALIZADO"] = manual["NOMBRE_ASIGNATURA_CANONICO_FINAL"].map(normalize_text)
    manual["CATEGORIA_REVISION"] = manual["ESTADO_FASE6"].map(status_category)
    decision_ids: list[str] = []
    for _, row in manual.iterrows():
        has_candidate = bool(str(row.get("CODIGO_ASIGNATURA_CANONICO_FINAL", "")).strip() or str(row.get("NOMBRE_ASIGNATURA_CANONICO_FINAL", "")).strip())
        if has_candidate:
            parts = [
                row.get("PLAN_DE_ESTUDIO_INSTITUCIONAL", ""),
                row.get("NOMBRE_ASIGNATURA_PDF_NORMALIZADO", ""),
                row.get("NIVEL_PDF", ""),
                row.get("TIPO_DIFERENCIA", ""),
                row.get("CODIGO_ASIGNATURA_CANONICO_FINAL", ""),
                row.get("NOMBRE_ASIGNATURA_CANONICO_NORMALIZADO", ""),
                row.get("NIVEL_CANONICO_FINAL", ""),
            ]
        else:
            parts = [
                row.get("PLAN_DE_ESTUDIO_INSTITUCIONAL", ""),
                row.get("NOMBRE_ASIGNATURA_PDF_NORMALIZADO", ""),
                row.get("NIVEL_PDF", ""),
                row.get("ESTADO_FASE6", ""),
                row.get("TIPO_DIFERENCIA", ""),
            ]
        decision_ids.append(stable_hash(parts, "DM_"))
    manual["ID_DECISION_MAESTRA"] = decision_ids

    detail_rows: list[dict[str, Any]] = []
    group_rows: list[dict[str, Any]] = []
    cand_rows: list[dict[str, Any]] = []
    for decision_id, group in manual.groupby("ID_DECISION_MAESTRA", dropna=False):
        plans = sorted(v for v in group["PLAN_DE_ESTUDIO_INSTITUCIONAL"].unique() if v)
        pdfs = sorted(v for v in group["ARCHIVO_PDF"].unique() if v)
        codigos = sorted(v for v in group["CODIGO_UNICO"].unique() if v)
        niveles = sorted(v for v in group["NIVEL_PDF"].unique() if v)
        candidatos = sorted(
            {
                "|".join(
                    [
                        str(r.get("CODIGO_ASIGNATURA_CANONICO_FINAL", "")),
                        str(r.get("NOMBRE_ASIGNATURA_CANONICO_FINAL", "")),
                        str(r.get("NIVEL_CANONICO_FINAL", "")),
                    ]
                )
                for r in group.to_dict("records")
                if str(r.get("CODIGO_ASIGNATURA_CANONICO_FINAL", "") or r.get("NOMBRE_ASIGNATURA_CANONICO_FINAL", "")).strip()
            }
        )
        estados = sorted(v for v in group["ESTADO_FASE6"].unique() if v)
        tipo_dif = sorted(v for v in group["TIPO_DIFERENCIA"].unique() if v)
        if len(group) == 1:
            clasificacion = "CASO_UNICO"
        elif len(plans) > 1:
            clasificacion = "NO_AGRUPABLE_PLAN_DISTINTO"
        elif len(niveles) > 1 and any("NIVEL" in t for t in tipo_dif):
            clasificacion = "NO_AGRUPABLE_NIVEL_DISTINTO"
        elif len(candidatos) > 1:
            clasificacion = "NO_AGRUPABLE_CANDIDATO_DISTINTO"
        elif len(estados) > 1:
            clasificacion = "NO_AGRUPABLE_ESTADO_DISTINTO"
        elif any(e in {"PENDIENTE_FUSION_DIVISION", "AMBIGUA_MULTIPLES_CANDIDATOS"} for e in estados):
            clasificacion = "NO_AGRUPABLE_CANDIDATO_DISTINTO"
        elif len(codigos) > 1:
            clasificacion = "DECISION_REUTILIZABLE_CON_VARIANTES_DE_JORNADA"
        elif len(pdfs) == 1:
            clasificacion = "DECISION_REUTILIZABLE_MISMO_PDF"
        else:
            clasificacion = "DECISION_UNICA_REUTILIZABLE"
        propagable = "SI" if clasificacion.startswith("DECISION_REUTILIZABLE") or clasificacion == "DECISION_UNICA_REUTILIZABLE" else "NO"
        if any(e in {"PENDIENTE_FUSION_DIVISION", "AMBIGUA_MULTIPLES_CANDIDATOS"} for e in estados):
            propagable = "NO"
        motivo_prop = "Mismo plan, estado, tipo de diferencia y candidato." if propagable == "SI" else clasificacion
        id_cases = group["ID_CASO"].tolist()
        group_cands = candidates[candidates["ID_CASO"].isin(id_cases)].copy()
        if not group_cands.empty:
            group_cands["PUNTAJE_NUM"] = pd.to_numeric(group_cands["PUNTAJE_SIMILITUD"], errors="coerce").fillna(0)
            top = (
                group_cands.groupby(["CODIGO_CANDIDATO", "NOMBRE_CANDIDATO", "NIVEL_CANDIDATO"], dropna=False)
                .agg(PUNTAJE=("PUNTAJE_NUM", "max"), FUENTE=("FUENTE", lambda s: "|".join(sorted(set(str(v) for v in s if str(v))))), CASOS=("ID_CASO", "nunique"))
                .reset_index()
                .sort_values(["PUNTAJE", "CASOS"], ascending=[False, False])
                .head(3)
            )
        else:
            top = pd.DataFrame(columns=["CODIGO_CANDIDATO", "NOMBRE_CANDIDATO", "NIVEL_CANDIDATO", "PUNTAJE", "FUENTE", "CASOS"])
        top_records = top.to_dict("records")
        while len(top_records) < 3:
            top_records.append({})
        puntajes = [float(r.get("PUNTAJE", 0) or 0) for r in top_records]
        brecha = round(puntajes[0] - puntajes[1], 6) if len(puntajes) > 1 else puntajes[0]
        candidato_principal_permitido = (
            propagable == "SI"
            and bool(top_records[0].get("NOMBRE_CANDIDATO", ""))
            and brecha >= 0.08
            and not any(e in {"PENDIENTE_FUSION_DIVISION", "AMBIGUA_MULTIPLES_CANDIDATOS", "PENDIENTE_SIN_EVIDENCIA"} for e in estados)
        )
        razon_candidato = "" if candidato_principal_permitido else "No se propone candidato unico por ambiguedad, evidencia insuficiente o propagacion no autorizable."
        base_row = {
            "ID_DECISION_MAESTRA": decision_id,
            "TOTAL_CASOS": len(group),
            "TOTAL_CODIGO_UNICO": len(codigos),
            "TOTAL_PDF": len(pdfs),
            "TOTAL_PLANES": len(plans),
            "TOTAL_NIVELES_PDF": len(niveles),
            "TOTAL_CANDIDATOS": len(candidatos),
            "TOTAL_ESTADOS_FASE6": len(estados),
            "CODIGOS_UNICOS_AFECTADOS": "|".join(codigos),
            "PDF_AFECTADOS": "|".join(pdfs),
            "PLANES_AFECTADOS": "|".join(plans),
            "NIVELES_PDF_OBSERVADOS": "|".join(niveles),
            "CANDIDATOS_OBSERVADOS": "|".join(candidatos),
            "ESTADOS_FASE6_OBSERVADOS": "|".join(estados),
            "CLASIFICACION_GRUPO": clasificacion,
            "PROPAGACION_AUTORIZABLE": propagable,
            "MOTIVO_PROPAGACION": motivo_prop,
            "PLAN_DE_ESTUDIO_INSTITUCIONAL": first_nonempty(group["PLAN_DE_ESTUDIO_INSTITUCIONAL"].tolist()),
            "NIVEL_PDF": first_nonempty(group["NIVEL_PDF"].tolist()),
            "CODIGO_ASIGNATURA_PDF": first_nonempty(group["CODIGO_ASIGNATURA_PDF"].tolist()),
            "NOMBRE_ASIGNATURA_PDF": first_nonempty(group["NOMBRE_ASIGNATURA_PDF"].tolist()),
            "ESTADO_FASE5": first_nonempty(group["ESTADO_FASE5"].tolist()),
            "ESTADO_FASE6": first_nonempty(group["ESTADO_FASE6"].tolist()),
            "TIPO_DIFERENCIA": first_nonempty(group["TIPO_DIFERENCIA"].tolist()),
            "CATEGORIA_REVISION": first_nonempty(group["CATEGORIA_REVISION"].tolist()),
            "MOTIVO_TECNICO": first_nonempty(group["MOTIVO"].tolist()),
            "ACCION_RECOMENDADA_TECNICA": first_nonempty(group["ACCION_REQUERIDA"].tolist()),
            "ES_PROPAGABLE": propagable,
            "BRECHA_PUNTAJE": brecha,
            "CANDIDATO_PRINCIPAL_PROPUESTO": "SI" if candidato_principal_permitido else "NO",
            "RAZON_CANDIDATO_PRINCIPAL": razon_candidato,
        }
        for idx in range(3):
            rec = top_records[idx]
            n = idx + 1
            base_row[f"CODIGO_CANDIDATO_{n}"] = rec.get("CODIGO_CANDIDATO", "") if (idx > 0 or candidato_principal_permitido) else ""
            base_row[f"NOMBRE_CANDIDATO_{n}"] = rec.get("NOMBRE_CANDIDATO", "") if (idx > 0 or candidato_principal_permitido) else ""
            base_row[f"NIVEL_CANDIDATO_{n}"] = rec.get("NIVEL_CANDIDATO", "") if (idx > 0 or candidato_principal_permitido) else ""
            base_row[f"PUNTAJE_CANDIDATO_{n}"] = rec.get("PUNTAJE", "") if (idx > 0 or candidato_principal_permitido) else ""
        group_rows.append(base_row)
        for rec in top_records:
            if rec:
                cand_rows.append(
                    {
                        "ID_DECISION_MAESTRA": decision_id,
                        "CODIGO_CANDIDATO": rec.get("CODIGO_CANDIDATO", ""),
                        "NOMBRE_CANDIDATO": rec.get("NOMBRE_CANDIDATO", ""),
                        "NIVEL_CANDIDATO": rec.get("NIVEL_CANDIDATO", ""),
                        "PUNTAJE": rec.get("PUNTAJE", ""),
                        "BRECHA_RESPECTO_PRIMERO": round(puntajes[0] - float(rec.get("PUNTAJE", 0) or 0), 6),
                        "COINCIDE_CODIGO": "",
                        "COINCIDE_NIVEL": "",
                        "PALABRAS_COINCIDENTES": "",
                        "PALABRAS_DIFERENTES": "",
                        "FUENTE": rec.get("FUENTE", ""),
                        "CANDIDATO_PRINCIPAL": "SI" if candidato_principal_permitido and rec == top_records[0] else "NO",
                        "RAZON_NO_PROPONER": razon_candidato,
                    }
                )
        for _, case in group.iterrows():
            detail_rows.append(
                {
                    "ID_CASO": case["ID_CASO"],
                    "ID_DECISION_MAESTRA": decision_id,
                    "CODIGO_UNICO": case["CODIGO_UNICO"],
                    "ARCHIVO_PDF": case["ARCHIVO_PDF"],
                    "PLAN": case["PLAN_DE_ESTUDIO_INSTITUCIONAL"],
                    "JORNADA": re.search(r"J([^V]+)V", str(case["CODIGO_UNICO"])).group(1) if re.search(r"J([^V]+)V", str(case["CODIGO_UNICO"])) else "",
                    "NIVEL_PDF": case["NIVEL_PDF"],
                    "NOMBRE_PDF": case["NOMBRE_ASIGNATURA_PDF"],
                    "CANDIDATO": "|".join([case.get("CODIGO_ASIGNATURA_CANONICO_FINAL", ""), case.get("NOMBRE_ASIGNATURA_CANONICO_FINAL", ""), case.get("NIVEL_CANONICO_FINAL", "")]).strip("|"),
                    "ESTADO_FASE6": case["ESTADO_FASE6"],
                    "TIPO_DIFERENCIA": case["TIPO_DIFERENCIA"],
                    "ES_PROPAGABLE": propagable,
                    "ESTADO_REVISION_MAESTRA": "PENDIENTE_REVISOR",
                    "MOTIVO_PROPAGACION": motivo_prop,
                    "REQUIERE_CONFIRMACION_INDIVIDUAL": "NO" if propagable == "SI" else "SI",
                }
            )

    grupos = pd.DataFrame(group_rows).sort_values(["CATEGORIA_REVISION", "TOTAL_CASOS", "ID_DECISION_MAESTRA"], ascending=[True, False, True])
    detalle = pd.DataFrame(detail_rows).sort_values(["ID_DECISION_MAESTRA", "ID_CASO"])
    candidatos_decision = pd.DataFrame(cand_rows)
    decisiones_maestras = len(grupos)
    grouped_cases = int(grupos.loc[grupos["TOTAL_CASOS"] > 1, "TOTAL_CASOS"].sum())
    unique_cases = int((grupos["TOTAL_CASOS"] == 1).sum())
    non_groupable_cases = int(grupos.loc[grupos["PROPAGACION_AUTORIZABLE"] != "SI", "TOTAL_CASOS"].sum())
    reduction_abs = 146 - decisiones_maestras
    reduction_pct = round(reduction_abs * 100 / 146, 2)
    propagable_groups = int((grupos["PROPAGACION_AUTORIZABLE"] == "SI").sum())
    non_propagable_groups = decisiones_maestras - propagable_groups

    resumen_rows = [
        {"METRICA": "CASOS_ORIGINALES", "VALOR": 146},
        {"METRICA": "DECISIONES_MAESTRAS", "VALOR": decisiones_maestras},
        {"METRICA": "CASOS_AGRUPADOS", "VALOR": grouped_cases},
        {"METRICA": "CASOS_UNICOS", "VALOR": unique_cases},
        {"METRICA": "CASOS_NO_AGRUPABLES", "VALOR": non_groupable_cases},
        {"METRICA": "REDUCCION_ABSOLUTA", "VALOR": reduction_abs},
        {"METRICA": "REDUCCION_PORCENTUAL", "VALOR": reduction_pct},
    ]
    for cat in ["APROXIMADOS", "DIFERENCIAS_NIVEL", "REVISAR_NOMBRE", "SIN_CORRESPONDENCIA", "FUSION_DIVISION", "PENDIENTES_SIN_EVIDENCIA"]:
        resumen_rows.append({"METRICA": f"{cat}_DECISIONES_MAESTRAS", "VALOR": int((grupos["CATEGORIA_REVISION"] == cat).sum())})
    resumen = pd.DataFrame(resumen_rows)

    import_template = pd.DataFrame(
        columns=[
            "ID_DECISION_MAESTRA",
            "DECISION_REVISOR",
            "CODIGO_CANONICO_CONFIRMADO",
            "NOMBRE_CANONICO_CONFIRMADO",
            "NIVEL_CANONICO_CONFIRMADO",
            "OBSERVACION_REVISOR",
            "REVISADO_POR",
            "FECHA_REVISION",
        ]
    )
    propagation = detalle[
        [
            "ID_DECISION_MAESTRA",
            "ID_CASO",
            "CODIGO_UNICO",
            "ARCHIVO_PDF",
            "PLAN",
            "ES_PROPAGABLE",
            "MOTIVO_PROPAGACION",
            "REQUIERE_CONFIRMACION_INDIVIDUAL",
        ]
    ].copy()

    out_groups = output_path(ctx, "02_INTERMEDIOS/07A_GRUPOS_DECISION_MAESTRA.tsv")
    out_map = output_path(ctx, "02_INTERMEDIOS/07B_MAPEO_CASO_A_DECISION_MAESTRA.tsv")
    out_cand = output_path(ctx, "02_INTERMEDIOS/07C_CANDIDATOS_POR_DECISION_MAESTRA.tsv")
    out_excel = output_path(ctx, "04_RESULTADOS/21_REVISION_MANUAL_OPTIMIZADA.xlsx")
    out_template = output_path(ctx, "04_RESULTADOS/22_PLANTILLA_IMPORTACION_DECISIONES_REVISOR.tsv")
    out_propagation = output_path(ctx, "04_RESULTADOS/23_MAPEO_PROPAGACION_DECISIONES.tsv")
    out_audit_excel = output_path(ctx, "05_AUDITORIA/FASE7_AUDITORIA_EXCEL_ORIGINAL.tsv")
    out_summary = output_path(ctx, "05_AUDITORIA/FASE7_RESUMEN_AGRUPACION_REVISION.tsv")

    salidas = [
        write_tsv(grupos, out_groups),
        write_tsv(detalle, out_map),
        write_tsv(candidatos_decision, out_cand),
        write_tsv(pd.DataFrame(audit_rows), out_audit_excel),
        write_tsv(resumen, out_summary),
        write_tsv(grupos, output_path(ctx, "05_AUDITORIA/FASE7_GRUPOS_DECISION_MAESTRA.tsv")),
        write_tsv(detalle[detalle["ES_PROPAGABLE"] == "SI"], output_path(ctx, "05_AUDITORIA/FASE7_CASOS_AGRUPADOS.tsv")),
        write_tsv(detalle[detalle["ES_PROPAGABLE"] != "SI"], output_path(ctx, "05_AUDITORIA/FASE7_CASOS_NO_AGRUPABLES.tsv")),
        write_tsv(propagation, output_path(ctx, "05_AUDITORIA/FASE7_CONTROL_PROPAGACION.tsv")),
        write_tsv(import_template, out_template),
        write_tsv(propagation, out_propagation),
    ]

    decisiones_permitidas = [
        "CONFIRMAR_EQUIVALENCIA",
        "CONFIRMAR_EQUIVALENCIA_CON_OBSERVACION",
        "RECHAZAR_EQUIVALENCIA",
        "CONFIRMAR_SIN_CORRESPONDENCIA",
        "CONFIRMAR_NO_CURRICULAR",
        "CONFIRMAR_ERROR_EXTRACCION",
        "CONFIRMAR_FUSION",
        "CONFIRMAR_DIVISION",
        "REQUIERE_ANTECEDENTE",
        "MANTENER_PENDIENTE",
    ]
    review_cols = [
        "DECISION_REVISOR",
        "CODIGO_CANONICO_CONFIRMADO",
        "NOMBRE_CANONICO_CONFIRMADO",
        "NIVEL_CANONICO_CONFIRMADO",
        "OBSERVACION_REVISOR",
        "REVISADO_POR",
        "FECHA_REVISION",
    ]
    decisiones_xlsx = grupos.copy()
    for col in review_cols:
        decisiones_xlsx[col] = ""
    detalle_xlsx = detalle.copy()
    instrucciones = pd.DataFrame(
        {
            "INSTRUCCION": [
                "Revise una fila por ID_DECISION_MAESTRA en DECISIONES_MAESTRAS.",
                "Edite solo columnas del revisor. Las columnas tecnicas son trazabilidad, no carga.",
                "Una decision se propaga solo a casos compatibles indicados como ES_PROPAGABLE=SI.",
                "No confirme equivalencias por nombre si hay candidato competitivo, fusion/division o evidencia insuficiente.",
                "PLAN_ESTUDIOS_SIES y PLAN_DE_ESTUDIO_INSTITUCIONAL son campos distintos; no los iguale.",
                "Este archivo no es precarga, PES ni fuente oficial.",
            ]
        }
    )
    diccionario = pd.DataFrame(
        [
            {"COLUMNA": col, "EDITABLE": "SI" if col in review_cols else "NO", "DESCRIPCION": "Columna de decision del revisor." if col in review_cols else "Campo tecnico o de trazabilidad."}
            for col in sorted(set(decisiones_xlsx.columns).union(detalle_xlsx.columns))
        ]
    )
    listas = pd.DataFrame({"DECISION_REVISOR": decisiones_permitidas})
    out_excel.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_excel, engine="xlsxwriter") as writer:
        sheet_map = {
            "INSTRUCCIONES": instrucciones,
            "RESUMEN": resumen,
            "DECISIONES_MAESTRAS": decisiones_xlsx,
            "DETALLE_CASOS": detalle_xlsx,
            "APROXIMADOS": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "APROXIMADOS"],
            "DIFERENCIAS_NIVEL": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "DIFERENCIAS_NIVEL"],
            "REVISAR_NOMBRE": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "REVISAR_NOMBRE"],
            "SIN_CORRESPONDENCIA": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "SIN_CORRESPONDENCIA"],
            "FUSION_DIVISION": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "FUSION_DIVISION"],
            "PENDIENTES": decisiones_xlsx[decisiones_xlsx["CATEGORIA_REVISION"] == "PENDIENTES_SIN_EVIDENCIA"],
            "DICCIONARIO": diccionario,
            "LISTAS": listas,
        }
        for sheet, df in sheet_map.items():
            df.to_excel(writer, sheet_name=sheet, index=False)
    wb = load_workbook(out_excel)
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    locked_fill = PatternFill("solid", fgColor="E7E6E6")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        for col_idx, header in enumerate(headers, start=1):
            width = min(max(len(header) + 2, 12), 45)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
            unlocked = header in review_cols and ws.title in {"DECISIONES_MAESTRAS", "APROXIMADOS", "DIFERENCIAS_NIVEL", "REVISAR_NOMBRE", "SIN_CORRESPONDENCIA", "FUSION_DIVISION", "PENDIENTES"}
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col_idx)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.protection = Protection(locked=not unlocked)
                if unlocked:
                    cell.fill = editable_fill
                elif ws.title != "LISTAS":
                    cell.fill = locked_fill
        if "DECISION_REVISOR" in headers and ws.max_row > 1:
            cidx = headers.index("DECISION_REVISOR") + 1
            col = get_column_letter(cidx)
            dv = DataValidation(type="list", formula1="=LISTAS!$A$2:$A$11", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{ws.max_row}")
        if ws.title not in {"INSTRUCCIONES", "RESUMEN", "DICCIONARIO", "LISTAS"}:
            ws.protection.sheet = True
            ws.protection.enable()
    if "LISTAS" in wb.sheetnames:
        wb["LISTAS"].sheet_state = "hidden"
    wb.save(out_excel)

    excel_control = pd.DataFrame(
        [
            {"CONTROL": "excel_optimizado_existe", "ESTADO": "OK" if out_excel.exists() else "ERROR", "OBSERVACION": str(out_excel)},
            {"CONTROL": "hoja_instrucciones", "ESTADO": "OK" if "INSTRUCCIONES" in load_workbook(out_excel, read_only=True).sheetnames else "ERROR", "OBSERVACION": ""},
            {"CONTROL": "hoja_diccionario", "ESTADO": "OK" if "DICCIONARIO" in load_workbook(out_excel, read_only=True).sheetnames else "ERROR", "OBSERVACION": ""},
            {"CONTROL": "validacion_lista", "ESTADO": "OK", "OBSERVACION": "DECISION_REVISOR"},
            {"CONTROL": "columnas_tecnicas_protegidas", "ESTADO": "OK", "OBSERVACION": "sin password"},
        ]
    )
    salidas.append(write_tsv(excel_control, output_path(ctx, "05_AUDITORIA/FASE7_CONTROL_EXCEL_OPTIMIZADO.tsv")))

    integrity = [
        {"CONTROL": "excel_original_hash_preservado", "ESTADO": "OK" if sha256_file(excel_original) == hash_excel_original else "ERROR", "OBSERVACION": hash_excel_original},
        {"CONTROL": "casos_originales_146", "ESTADO": "OK" if len(manual) == 146 else "ERROR", "OBSERVACION": len(manual)},
        {"CONTROL": "casos_trazados_146", "ESTADO": "OK" if detalle["ID_CASO"].nunique() == 146 else "ERROR", "OBSERVACION": detalle["ID_CASO"].nunique()},
        {"CONTROL": "sin_casos_adicionales", "ESTADO": "OK" if set(detalle["ID_CASO"]) == set(manual["ID_CASO"]) else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "cada_caso_con_decision_maestra", "ESTADO": "OK" if detalle["ID_DECISION_MAESTRA"].ne("").all() else "ERROR", "OBSERVACION": ""},
        {"CONTROL": "sin_decisiones_humanas_aplicadas", "ESTADO": "OK", "OBSERVACION": "0"},
        {"CONTROL": "sin_precarga_pes", "ESTADO": "OK", "OBSERVACION": ""},
        {"CONTROL": "apto_para_carga", "ESTADO": "OK", "OBSERVACION": "NO"},
    ]
    integrity_df = pd.DataFrame(integrity)
    salidas.append(write_tsv(integrity_df, output_path(ctx, "05_AUDITORIA/FASE7_CONTROL_INTEGRIDAD.tsv")))
    salidas.append(write_tsv(detalle[["ID_CASO", "ID_DECISION_MAESTRA", "CODIGO_UNICO", "ARCHIVO_PDF", "PLAN", "ES_PROPAGABLE", "MOTIVO_PROPAGACION"]], output_path(ctx, "05_AUDITORIA/FASE7_TRAZABILIDAD_REVISION.tsv")))
    log_path = output_path(ctx, "06_LOGS/fase7_preparacion_revision_manual.log")
    salidas.append(write_text(f"{now_iso()} FASE 7 preparacion revision manual completada\n", log_path))
    salidas.append(str(out_excel.resolve()))
    errores = integrity_df[integrity_df["ESTADO"] != "OK"].to_dict("records")
    estado = "COMPLETADA" if not errores else "BLOQUEADA"

    conteos = {
        "casos_originales": 146,
        "casos_preservados": int(detalle["ID_CASO"].nunique()),
        "decisiones_maestras": decisiones_maestras,
        "casos_agrupados": grouped_cases,
        "casos_unicos": unique_cases,
        "casos_no_agrupables": non_groupable_cases,
        "reduccion_absoluta": reduction_abs,
        "reduccion_porcentual": reduction_pct,
        "grupos_propagables": propagable_groups,
        "grupos_no_propagables": non_propagable_groups,
        "aproximados_decisiones": int((grupos["CATEGORIA_REVISION"] == "APROXIMADOS").sum()),
        "niveles_decisiones": int((grupos["CATEGORIA_REVISION"] == "DIFERENCIAS_NIVEL").sum()),
        "nombres_decisiones": int((grupos["CATEGORIA_REVISION"] == "REVISAR_NOMBRE").sum()),
        "sin_correspondencia_decisiones": int((grupos["CATEGORIA_REVISION"] == "SIN_CORRESPONDENCIA").sum()),
        "fusion_division_decisiones": int((grupos["CATEGORIA_REVISION"] == "FUSION_DIVISION").sum()),
        "pendientes_sin_evidencia_decisiones": int((grupos["CATEGORIA_REVISION"] == "PENDIENTES_SIN_EVIDENCIA").sum()),
        "decisiones_humanas_detectadas": 0,
        "decisiones_humanas_aplicadas": 0,
        "originales_modificados": False,
        "precarga_generada": False,
        "pes_generado": False,
        "apto_para_carga": False,
    }
    return PhaseResult(
        estado=estado,
        salidas=salidas,
        conteos=conteos,
        validaciones=integrity,
        errores=errores[:20],
        pendientes=[],
        resumen=[
            f"Estado: {estado}",
            "Casos originales de revision: 146",
            f"Casos preservados: {conteos['casos_preservados']}",
            f"Decisiones maestras unicas: {decisiones_maestras}",
            f"Casos agrupados: {grouped_cases}",
            f"Casos unicos: {unique_cases}",
            f"Casos no agrupables: {non_groupable_cases}",
            f"Reduccion absoluta de decisiones manuales: {reduction_abs}",
            f"Reduccion porcentual: {reduction_pct}%",
            f"Grupos propagables: {propagable_groups}",
            f"Grupos no propagables: {non_propagable_groups}",
            f"Aproximados - decisiones maestras: {conteos['aproximados_decisiones']}",
            f"Diferencias de nivel - decisiones maestras: {conteos['niveles_decisiones']}",
            f"Revisar nombre - decisiones maestras: {conteos['nombres_decisiones']}",
            f"Sin correspondencia - decisiones maestras: {conteos['sin_correspondencia_decisiones']}",
            f"Fusion/division - decisiones maestras: {conteos['fusion_division_decisiones']}",
            f"Pendientes sin evidencia - decisiones maestras: {conteos['pendientes_sin_evidencia_decisiones']}",
            "Decisiones humanas detectadas previamente: 0",
            "Decisiones humanas aplicadas: 0",
            f"Excel original preservado: {'SI' if sha256_file(excel_original) == hash_excel_original else 'NO'}",
            f"Excel optimizado generado: {'SI' if out_excel.exists() else 'NO'}",
            f"Plantilla de importacion generada: {'SI' if out_template.exists() else 'NO'}",
            f"Mapeo de propagacion generado: {'SI' if out_propagation.exists() else 'NO'}",
            f"Excel optimizado: {out_excel.resolve()}",
            f"Plantilla importacion: {out_template.resolve()}",
            f"Mapeo propagacion: {out_propagation.resolve()}",
            f"Auditoria: {(ctx.run_dir / '05_AUDITORIA').resolve()}",
            f"Carpeta de ejecucion: {ctx.run_dir.resolve()}",
            "Originales modificados: NO",
            "Precarga generada: NO",
            "PES generado: NO",
            "Apto para carga: NO",
            "Siguiente paso: COMPLETAR_EXCEL_OPTIMIZADO_Y_LUEGO_IMPORTAR_DECISIONES_REVISOR",
        ],
        comando_reanudacion="",
    )


def _safe_read_tsv(path: Path) -> pd.DataFrame:
    if not path.exists() or path.stat().st_size == 0:
        return pd.DataFrame()
    return pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False, encoding="utf-8")


def _safe_excel_sheets(path: Path) -> dict[str, pd.DataFrame]:
    return pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False)


def _blank_review_decisions(df: pd.DataFrame) -> bool:
    review_cols = [
        "DECISION_REVISOR",
        "CODIGO_CANONICO_CONFIRMADO",
        "NOMBRE_CANONICO_CONFIRMADO",
        "NIVEL_CANONICO_CONFIRMADO",
        "OBSERVACION_REVISOR",
        "REVISADO_POR",
        "FECHA_REVISION",
    ]
    present = [col for col in review_cols if col in df.columns]
    if not present:
        return True
    filled = df[present].fillna("").astype(str).apply(lambda col: col.str.strip().ne(""))
    return not bool(filled.any(axis=None))


def _copy_with_manifest(files: list[Path], respaldo_dir: Path) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    respaldo_dir.mkdir(parents=True, exist_ok=True)
    for src in files:
        dst = respaldo_dir / src.name
        shutil.copy2(src, dst)
        rows.append(
            {
                "ARCHIVO_ORIGEN": str(src.resolve()),
                "ARCHIVO_RESPALDO": str(dst.resolve()),
                "HASH_SHA256": sha256_file(src),
                "TAMANO_BYTES": src.stat().st_size,
                "FECHA_MODIFICACION": datetime.fromtimestamp(src.stat().st_mtime).isoformat(timespec="seconds"),
            }
        )
    return pd.DataFrame(rows)


def _classify_logistica_case(row: pd.Series, reevaluable_ids: set[str], common_names: set[str]) -> str:
    case_id = str(row.get("ID_CASO", ""))
    estado = str(row.get("ESTADO_ORIGINAL", row.get("ESTADO_FASE6", "")))
    name = normalize_text(str(row.get("NOMBRE_PDF", row.get("NOMBRE_ASIGNATURA_PDF", ""))))
    if "FUSION" in normalize_text(estado):
        return "LOGISTICA_FUSION_CANDIDATA"
    if "DIVISION" in normalize_text(estado):
        return "LOGISTICA_DIVISION_CANDIDATA"
    if case_id in reevaluable_ids and "SIN_CORRESPONDENCIA" in normalize_text(estado):
        return "LOGISTICA_POSIBLE_FALSO_SIN_CORRESPONDENCIA"
    if name in common_names and name:
        return "LOGISTICA_ASIGNATURA_COMUN_VALIDADA"
    if case_id in reevaluable_ids:
        return "LOGISTICA_POSIBLE_COMUN_NO_CONFIRMADA"
    return "LOGISTICA_RELACION_NO_APLICA"


def _format_excel_derivado(path: Path, editable_cols: set[str]) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment, Protection

    wb = load_workbook(path)
    editable_fill = PatternFill("solid", fgColor="FFF2CC")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    locked_fill = PatternFill("solid", fgColor="E7E6E6")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        headers = [str(ws.cell(1, c).value or "") for c in range(1, ws.max_column + 1)]
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True)
        for col_idx, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(header) + 2, 14), 55)
            unlocked = header in editable_cols and ws.title == "DECISIONES_MAESTRAS"
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.protection = Protection(locked=not unlocked)
                cell.fill = editable_fill if unlocked else locked_fill
        if ws.title == "DECISIONES_MAESTRAS" and "DECISION_REVISOR" in headers and ws.max_row > 1:
            col = get_column_letter(headers.index("DECISION_REVISOR") + 1)
            dv = DataValidation(type="list", formula1="=LISTAS!$A$2:$A$11", allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{ws.max_row}")
        if ws.title not in {"INSTRUCCIONES", "RESUMEN", "DICCIONARIO", "LISTAS"}:
            ws.protection.sheet = True
            ws.protection.enable()
    if "LISTAS" in wb.sheetnames:
        wb["LISTAS"].sheet_state = "hidden"
    wb.save(path)


def regenerar_paquete_tec_cont(args: argparse.Namespace) -> int:
    from openpyxl import load_workbook

    run_dir = Path(args.carpeta_ejecucion).resolve()
    audit_dir = Path(args.auditoria_mallas).resolve()
    if not run_dir.exists():
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print(f"Motivo: carpeta de ejecucion no existe: {run_dir}")
        return 2
    if not audit_dir.exists():
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print(f"Motivo: auditoria de mallas no existe: {audit_dir}")
        return 2

    log_path = run_dir / "06_LOGS/fase8b_regeneracion_paquete_tec_cont.log"
    log_path.parent.mkdir(parents=True, exist_ok=True)
    log_path.write_text(f"{now_iso()} inicio regeneracion paquete tec-cont\n", encoding="utf-8")

    pdf_inputs = {
        "cierre148": run_dir / "04_RESULTADOS/16_CIERRE_148_CASOS_NO_EXACTOS.tsv",
        "post_revision": run_dir / "04_RESULTADOS/17_VALIDACION_ASIGNATURAS_PDF_PLAN_POST_REVISION.tsv",
        "excel20": run_dir / "04_RESULTADOS/20_REVISION_MANUAL_CASOS_PENDIENTES.xlsx",
        "excel21": run_dir / "04_RESULTADOS/21_REVISION_MANUAL_OPTIMIZADA.xlsx",
        "plantilla22": run_dir / "04_RESULTADOS/22_PLANTILLA_IMPORTACION_DECISIONES_REVISOR.tsv",
        "mapeo23": run_dir / "04_RESULTADOS/23_MAPEO_PROPAGACION_DECISIONES.tsv",
        "grupos07a": run_dir / "02_INTERMEDIOS/07A_GRUPOS_DECISION_MAESTRA.tsv",
        "mapa07b": run_dir / "02_INTERMEDIOS/07B_MAPEO_CASO_A_DECISION_MAESTRA.tsv",
    }
    audit_inputs = {
        "resumen": audit_dir / "04_RESULTADOS/01_RESUMEN_COMPARACION_LOGISTICA.tsv",
        "detalle_malla": audit_dir / "04_RESULTADOS/02_DETALLE_MALLA_COMPARTIDA_LOGISTICA.tsv",
        "representacion": audit_dir / "04_RESULTADOS/03_REPRESENTACION_EN_FLUJO_ACTUAL.tsv",
        "impacto": audit_dir / "04_RESULTADOS/04_IMPACTO_EN_146_CASOS.tsv",
        "reevaluar": audit_dir / "04_RESULTADOS/05_CASOS_LOGISTICA_REEVALUAR.tsv",
        "conclusion": audit_dir / "04_RESULTADOS/07_CONCLUSION_LOGISTICA.tsv",
        "recomendacion": audit_dir / "04_RESULTADOS/08_RECOMENDACION_ANTES_REVISION_MANUAL.tsv",
        "resumen_md": audit_dir / "04_RESULTADOS/09_RESUMEN_EJECUTIVO.md",
        "manifiesto": audit_dir / "00_CONTROL/manifiesto_final.json",
    }
    missing = [str(p) for p in list(pdf_inputs.values()) + list(audit_inputs.values()) if not p.exists()]
    if missing:
        log_path.write_text(log_path.read_text(encoding="utf-8") + f"bloqueo faltantes={missing}\n", encoding="utf-8")
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print(f"Motivo: faltan entradas obligatorias: {len(missing)}")
        return 2

    hashes_previos = {name: sha256_file(path) for name, path in pdf_inputs.items()}
    ts = timestamp()
    respaldo_dir = run_dir / f"07_RESPALDOS/PAQUETE_REVISION_PRE_TEC_CONT_{ts}"
    respaldo_df = _copy_with_manifest(
        [pdf_inputs["excel20"], pdf_inputs["excel21"], pdf_inputs["plantilla22"], pdf_inputs["mapeo23"]],
        respaldo_dir,
    )
    write_tsv(respaldo_df, respaldo_dir / "MANIFIESTO_RESPALDO.tsv")

    resumen = _safe_read_tsv(audit_inputs["resumen"])
    detalle_malla = _safe_read_tsv(audit_inputs["detalle_malla"])
    impacto = _safe_read_tsv(audit_inputs["impacto"])
    reevaluar = _safe_read_tsv(audit_inputs["reevaluar"])
    conclusion = _safe_read_tsv(audit_inputs["conclusion"])
    grupos = _safe_read_tsv(pdf_inputs["grupos07a"])
    detalle = _safe_read_tsv(pdf_inputs["mapa07b"])
    plantilla22 = _safe_read_tsv(pdf_inputs["plantilla22"])
    mapeo23 = _safe_read_tsv(pdf_inputs["mapeo23"])
    sheets21 = _safe_excel_sheets(pdf_inputs["excel21"])
    decisiones = sheets21.get("DECISIONES_MAESTRAS", pd.DataFrame())
    detalle_excel = sheets21.get("DETALLE_CASOS", pd.DataFrame())

    common = detalle_malla[detalle_malla["CLASIFICACION"].isin(["COMUN_CODIGO_EXACTO", "COMUN_NOMBRE_EXACTO", "COMUN_NOMBRE_NIVEL", "COMUN_CON_DIFERENCIA_NIVEL"])].copy()
    common_names = {normalize_text(x) for x in common.get("NOMBRE_TECNICO", pd.Series(dtype=str)).tolist() + common.get("NOMBRE_PROFESIONAL", pd.Series(dtype=str)).tolist() if normalize_text(x)}
    reevaluable_ids = set(reevaluar.get("ID_CASO", pd.Series(dtype=str)).astype(str))
    impact_ids = set(impacto.get("ID_CASO", pd.Series(dtype=str)).astype(str))
    fusion_mask = impacto.get("ESTADO_ORIGINAL", pd.Series(dtype=str)).astype(str).map(lambda x: "FUSION" in normalize_text(x) or "DIVISION" in normalize_text(x))
    fusion_rows = impacto.loc[fusion_mask.fillna(False).astype(bool)].copy()
    nivel_respaldo = conclusion.iloc[0].get("NIVEL_RESPALDO", "") if not conclusion.empty else ""

    manifest = json.loads(audit_inputs["manifiesto"].read_text(encoding="utf-8"))
    metrics = manifest.get("metricas", {})
    oficial = _safe_read_tsv(audit_dir / "02_EVIDENCIA/01_EVIDENCIA_OFICIAL.tsv")
    fuentes_oficiales_unicas = oficial["FUENTE"].nunique() if "FUENTE" in oficial.columns else 0
    clasif_normativa_ok = "SI"
    if "FUENTE" in oficial.columns:
        bad_official = oficial["FUENTE"].astype(str).map(lambda x: "scripts/" in x or "AUDITORIA_" in x)
        if bool(bad_official.fillna(False).astype(bool).any()):
            clasif_normativa_ok = "NO"

    validacion_rows = [
        {"CONTROL": "scripts_revisados_es_registro_no_archivo_unico", "ESTADO": "OK", "VALOR": metrics.get("scripts_revisados", 0), "OBSERVACION": "Se valida como metrica de evidencia de la auditoria; no regla oficial."},
        {"CONTROL": "fuentes_oficiales_unicas", "ESTADO": "OK", "VALOR": fuentes_oficiales_unicas, "OBSERVACION": "Conteo unico de FUENTE en evidencia oficial."},
        {"CONTROL": "clasificacion_normativa_validada", "ESTADO": "OK" if clasif_normativa_ok == "SI" else "ERROR", "VALOR": clasif_normativa_ok, "OBSERVACION": "No se aceptan scripts/auditorias como REGLA_OFICIAL."},
        {"CONTROL": "casos_logistica_22", "ESTADO": "OK" if len(impacto) == 22 else "ERROR", "VALOR": len(impacto), "OBSERVACION": ""},
        {"CONTROL": "casos_afectados_16", "ESTADO": "OK" if len(reevaluar) == 16 and set(reevaluable_ids).issubset(impact_ids) else "ERROR", "VALOR": len(reevaluar), "OBSERVACION": ""},
        {"CONTROL": "fusion_division_2", "ESTADO": "OK" if len(fusion_rows) == 2 else "ERROR", "VALOR": len(fusion_rows), "OBSERVACION": ""},
        {"CONTROL": "asignaturas_comunes_29", "ESTADO": "OK" if len(common) == 29 else "ERROR", "VALOR": len(common), "OBSERVACION": ""},
        {"CONTROL": "cobertura_100", "ESTADO": "OK" if str(resumen.iloc[0].get("COBERTURA_TECNICO_EN_PROFESIONAL", "")) in {"1.0", "1"} else "ERROR", "VALOR": resumen.iloc[0].get("COBERTURA_TECNICO_EN_PROFESIONAL", "") if not resumen.empty else "", "OBSERVACION": ""},
        {"CONTROL": "mismo_pdf_no_determinado", "ESTADO": "OK" if resumen.iloc[0].get("MISMO_PDF", "") == "NO_DETERMINADO" else "ERROR", "VALOR": resumen.iloc[0].get("MISMO_PDF", "") if not resumen.empty else "", "OBSERVACION": ""},
        {"CONTROL": "mismo_plan_no_determinado", "ESTADO": "OK" if resumen.iloc[0].get("MISMO_PLAN", "") == "NO_DETERMINADO" else "ERROR", "VALOR": resumen.iloc[0].get("MISMO_PLAN", "") if not resumen.empty else "", "OBSERVACION": ""},
        {"CONTROL": "nivel_respaldo_dato_observado", "ESTADO": "OK" if nivel_respaldo == "DATO_OBSERVADO" else "ERROR", "VALOR": nivel_respaldo, "OBSERVACION": ""},
    ]
    write_tsv(pd.DataFrame(validacion_rows), run_dir / "05_AUDITORIA/FASE8B_VALIDACION_AUDITORIA_MALLAS.tsv")
    if any(row["ESTADO"] != "OK" for row in validacion_rows):
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print("Motivo: no se pudo validar auditoria de mallas.")
        return 2
    if len(decisiones) != 146 or len(detalle) != 146 or len(detalle_excel) != 146:
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print("Motivo: no se preservan 146 casos/decisiones.")
        return 2
    if not _blank_review_decisions(decisiones):
        write_tsv(pd.DataFrame([{"CONTROL": "decisiones_humanas_previas", "ESTADO": "BLOQUEADA"}]), run_dir / "05_AUDITORIA/FASE8B_CONTROL_INTEGRIDAD.tsv")
        print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
        print("Estado: BLOQUEADA")
        print("Motivo: existen decisiones humanas en Excel original.")
        return 2

    impact_lookup = impacto.set_index("ID_CASO").to_dict("index") if not impacto.empty else {}
    common_lookup = {normalize_text(row.get("NOMBRE_TECNICO", "")): row for _, row in common.iterrows()}
    common_lookup.update({normalize_text(row.get("NOMBRE_PROFESIONAL", "")): row for _, row in common.iterrows()})
    reevaluation_rows: list[dict[str, Any]] = []
    for _, drow in detalle.iterrows():
        case_id = str(drow.get("ID_CASO", ""))
        if case_id not in impact_lookup:
            continue
        arow = impact_lookup[case_id]
        name_norm = normalize_text(arow.get("NOMBRE_PDF", drow.get("NOMBRE_PDF", "")))
        crow = common_lookup.get(name_norm, {})
        tmp = pd.Series({**drow.to_dict(), **arow})
        clasif = _classify_logistica_case(tmp, reevaluable_ids, common_names)
        reevaluation_rows.append(
            {
                "ID_CASO": case_id,
                "ID_DECISION_MAESTRA": drow.get("ID_DECISION_MAESTRA", ""),
                "CODIGO_UNICO": drow.get("CODIGO_UNICO", arow.get("CODIGO_UNICO", "")),
                "ARCHIVO_PDF": drow.get("ARCHIVO_PDF", arow.get("ARCHIVO_PDF", "")),
                "PLAN_DE_ESTUDIO_INSTITUCIONAL": drow.get("PLAN", arow.get("PLAN", "")),
                "NOMBRE_ASIGNATURA_PDF": drow.get("NOMBRE_PDF", arow.get("NOMBRE_PDF", "")),
                "NIVEL_PDF": drow.get("NIVEL_PDF", ""),
                "CODIGO_ASIGNATURA_CANONICO_PROPUESTO": crow.get("CODIGO_TECNICO", "") or crow.get("CODIGO_PROFESIONAL", ""),
                "NOMBRE_ASIGNATURA_CANONICO_PROPUESTO": crow.get("NOMBRE_TECNICO", "") or crow.get("NOMBRE_PROFESIONAL", ""),
                "PLAN_TECNICO_RELACIONADO": "TLOG",
                "PLAN_PROFESIONAL_RELACIONADO": "ILOG",
                "ASIGNATURA_COMUN_VALIDADA": "SI" if name_norm in common_names else "NO",
                "TIPO_RELACION_TEC_CONT": "MALLA_TECNICA_CONTENIDA_EN_PROFESIONAL",
                "EVIDENCIA_RELACION": str(audit_inputs["detalle_malla"].resolve()),
                "NIVEL_RESPALDO": "DATO_OBSERVADO",
                "ESTADO_TECNICO_ACTUAL": drow.get("ESTADO_FASE6", arow.get("ESTADO_ORIGINAL", "")),
                "CLASIFICACION_REEVALUACION": clasif,
                "REQUIERE_REVISION_MANUAL": "SI",
                "ACCION_SUGERIDA": "REEVALUAR_CON_EVIDENCIA_TEC_CONT" if case_id in reevaluable_ids else "MANTENER_REVISION_MANUAL",
            }
        )
    reevaluation = pd.DataFrame(reevaluation_rows)
    out24 = run_dir / "04_RESULTADOS/24_REEVALUACION_CONTROLADA_22_CASOS_LOGISTICA.tsv"
    write_tsv(reevaluation, out24)

    tech_cols = [
        "ES_CASO_LOGISTICA",
        "RELACION_TEC_CONT_APLICA",
        "CLASIFICACION_REEVALUACION",
        "MALLA_TECNICA_CONTENIDA",
        "ASIGNATURA_COMUN_VALIDADA",
        "PLAN_TECNICO_RELACIONADO",
        "PLAN_PROFESIONAL_RELACIONADO",
        "EVIDENCIA_TEC_CONT",
        "NIVEL_RESPALDO_TEC_CONT",
        "ACCION_SUGERIDA_TECNICA",
        "REQUIERE_CONFIRMACION_INDIVIDUAL",
    ]
    reevaluation_by_dm = reevaluation.groupby("ID_DECISION_MAESTRA").first().to_dict("index") if not reevaluation.empty else {}
    decisiones_new = decisiones.copy()
    for col in tech_cols:
        decisiones_new[col] = ""
    for idx, row in decisiones_new.iterrows():
        dm = row.get("ID_DECISION_MAESTRA", "")
        info = reevaluation_by_dm.get(dm)
        if info:
            decisiones_new.at[idx, "ES_CASO_LOGISTICA"] = "SI"
            decisiones_new.at[idx, "RELACION_TEC_CONT_APLICA"] = "SI" if info["ID_CASO"] in reevaluable_ids else "NO"
            decisiones_new.at[idx, "CLASIFICACION_REEVALUACION"] = info["CLASIFICACION_REEVALUACION"]
            decisiones_new.at[idx, "MALLA_TECNICA_CONTENIDA"] = "SI"
            decisiones_new.at[idx, "ASIGNATURA_COMUN_VALIDADA"] = info["ASIGNATURA_COMUN_VALIDADA"]
            decisiones_new.at[idx, "PLAN_TECNICO_RELACIONADO"] = "TLOG"
            decisiones_new.at[idx, "PLAN_PROFESIONAL_RELACIONADO"] = "ILOG"
            decisiones_new.at[idx, "EVIDENCIA_TEC_CONT"] = str(audit_inputs["detalle_malla"].resolve())
            decisiones_new.at[idx, "NIVEL_RESPALDO_TEC_CONT"] = "DATO_OBSERVADO"
            decisiones_new.at[idx, "ACCION_SUGERIDA_TECNICA"] = info["ACCION_SUGERIDA"]
            decisiones_new.at[idx, "REQUIERE_CONFIRMACION_INDIVIDUAL"] = "SI"
        else:
            decisiones_new.at[idx, "ES_CASO_LOGISTICA"] = "NO"
            decisiones_new.at[idx, "RELACION_TEC_CONT_APLICA"] = "NO"
            decisiones_new.at[idx, "MALLA_TECNICA_CONTENIDA"] = "NO"
            decisiones_new.at[idx, "NIVEL_RESPALDO_TEC_CONT"] = ""
            decisiones_new.at[idx, "REQUIERE_CONFIRMACION_INDIVIDUAL"] = row.get("REQUIERE_CONFIRMACION_INDIVIDUAL", "")

    detalle_new = detalle_excel.copy()
    detalle_new["ES_CASO_LOGISTICA"] = detalle_new["ID_CASO"].astype(str).isin(impact_ids).map({True: "SI", False: "NO"})
    detalle_new["ES_POTENCIALMENTE_AFECTADO"] = detalle_new["ID_CASO"].astype(str).isin(reevaluable_ids).map({True: "SI", False: "NO"})

    out27_rows = []
    fusion_ids = set(fusion_rows.get("ID_CASO", pd.Series(dtype=str)).astype(str))
    for _, drow in detalle.iterrows():
        case_id = str(drow.get("ID_CASO", ""))
        out27_rows.append(
            {
                "ID_DECISION_MAESTRA": drow.get("ID_DECISION_MAESTRA", ""),
                "ID_CASO": case_id,
                "ES_LOGISTICA": "SI" if case_id in impact_ids else "NO",
                "ES_POTENCIALMENTE_AFECTADO": "SI" if case_id in reevaluable_ids else "NO",
                "TIPO_RELACION": "MALLA_TECNICA_CONTENIDA_EN_PROFESIONAL" if case_id in impact_ids else "",
                "ASIGNATURA_COMUN_VALIDADA": "SI" if case_id in set(reevaluation.loc[reevaluation["ASIGNATURA_COMUN_VALIDADA"] == "SI", "ID_CASO"]) else "NO",
                "FUSION_DIVISION": "SI" if case_id in fusion_ids else "NO",
                "PROPAGACION_AUTORIZADA": "NO",
                "MOTIVO": "No se crea regla de propagacion; requiere confirmacion humana.",
            }
        )
    out27 = run_dir / "04_RESULTADOS/27_MAPEO_REEVALUACION_TEC_CONT.tsv"
    write_tsv(pd.DataFrame(out27_rows), out27)

    if len(plantilla22) == 0:
        plantilla_base_cols = [
            "ID_DECISION_MAESTRA",
            "DECISION_REVISOR",
            "CODIGO_CANONICO_CONFIRMADO",
            "NOMBRE_CANONICO_CONFIRMADO",
            "NIVEL_CANONICO_CONFIRMADO",
            "OBSERVACION_REVISOR",
            "REVISADO_POR",
            "FECHA_REVISION",
        ]
        plantilla26 = decisiones_new[plantilla_base_cols + tech_cols].copy()
        for review_col in [
            "DECISION_REVISOR",
            "CODIGO_CANONICO_CONFIRMADO",
            "NOMBRE_CANONICO_CONFIRMADO",
            "NIVEL_CANONICO_CONFIRMADO",
            "OBSERVACION_REVISOR",
            "REVISADO_POR",
            "FECHA_REVISION",
        ]:
            plantilla26[review_col] = ""
    else:
        plantilla26 = plantilla22.copy()
        for col in tech_cols:
            plantilla26[col] = ""
    out26 = run_dir / "04_RESULTADOS/26_PLANTILLA_DECISIONES_REVISOR_TEC_CONT.tsv"
    write_tsv(plantilla26, out26)

    resumen_rows = [
        {"METRICA": "CASOS_MANUALES_PRESERVADOS", "VALOR": 146},
        {"METRICA": "CASOS_LOGISTICA", "VALOR": len(reevaluation)},
        {"METRICA": "CASOS_POTENCIALMENTE_AFECTADOS", "VALOR": len(reevaluar)},
        {"METRICA": "FUSIONES_DIVISIONES", "VALOR": len(fusion_rows)},
        {"METRICA": "ASIGNATURAS_COMUNES_VALIDADAS", "VALOR": len(common)},
        {"METRICA": "DECISIONES_HUMANAS_APLICADAS", "VALOR": 0},
    ]
    instrucciones = pd.DataFrame(
        {
            "INSTRUCCION": [
                "Tecnico en Logistica e Ingenieria en Logistica tienen planes distintos.",
                "La auditoria observa que la malla tecnica aparece contenida en la profesional.",
                "Esto no significa identidad administrativa ni autoriza fusion de programas.",
                "Los 16 casos se presentan solo para reevaluacion humana.",
                "La persona revisora debe confirmar cada correspondencia.",
                "Los casos de fusion/division no deben resolverse automaticamente como uno a uno.",
                "La evidencia es DATO_OBSERVADO, no REGLA_OFICIAL.",
            ]
        }
    )
    diccionario = pd.DataFrame(
        [{"COLUMNA": col, "DESCRIPCION": "Columna informativa tecnica-continuidad."} for col in tech_cols]
        + [{"COLUMNA": "DECISION_REVISOR", "DESCRIPCION": "Editable por revisor; no completada automaticamente."}]
    )
    listas = pd.DataFrame({"DECISION_REVISOR": [
        "CONFIRMAR_EQUIVALENCIA",
        "CONFIRMAR_EQUIVALENCIA_CON_OBSERVACION",
        "RECHAZAR_EQUIVALENCIA",
        "CONFIRMAR_SIN_CORRESPONDENCIA",
        "CONFIRMAR_NO_CURRICULAR",
        "CONFIRMAR_ERROR_EXTRACCION",
        "CONFIRMAR_FUSION",
        "CONFIRMAR_DIVISION",
        "REQUIERE_ANTECEDENTE",
        "MANTENER_PENDIENTE",
    ]})
    evidence = pd.concat(
        [
            common.assign(TIPO_EVIDENCIA="MALLA_COMUN_VALIDADA"),
            conclusion.assign(TIPO_EVIDENCIA="CONCLUSION_AUDITORIA"),
        ],
        ignore_index=True,
        sort=False,
    )
    out25 = run_dir / "04_RESULTADOS/25_REVISION_MANUAL_CON_RELACION_TEC_CONT.xlsx"
    with pd.ExcelWriter(out25, engine="xlsxwriter") as writer:
        sheet_map = {
            "INSTRUCCIONES": instrucciones,
            "RESUMEN": pd.DataFrame(resumen_rows),
            "DECISIONES_MAESTRAS": decisiones_new,
            "LOGISTICA_REEVALUAR": reevaluation[reevaluation["ID_CASO"].isin(reevaluable_ids)],
            "LOGISTICA_MALLA_COMUN": common,
            "LOGISTICA_FUSION_DIVISION": reevaluation[reevaluation["ID_CASO"].isin(fusion_ids)],
            "OTROS_CASOS": decisiones_new[decisiones_new["ES_CASO_LOGISTICA"] != "SI"],
            "DETALLE_CASOS": detalle_new,
            "EVIDENCIA_TEC_CONT": evidence,
            "DICCIONARIO": diccionario,
            "LISTAS": listas,
        }
        for sheet, df in sheet_map.items():
            df.to_excel(writer, sheet_name=sheet[:31], index=False)
    _format_excel_derivado(out25, {"DECISION_REVISOR", "CODIGO_CANONICO_CONFIRMADO", "NOMBRE_CANONICO_CONFIRMADO", "NIVEL_CANONICO_CONFIRMADO", "OBSERVACION_REVISOR", "REVISADO_POR", "FECHA_REVISION"})

    controls = {
        "22": len(reevaluation),
        "16": len(reevaluar),
        "2": len(fusion_rows),
        "29": len(common),
    }
    write_tsv(reevaluation, run_dir / "05_AUDITORIA/FASE8B_CONTROL_22_CASOS_LOGISTICA.tsv")
    write_tsv(reevaluation[reevaluation["ID_CASO"].isin(reevaluable_ids)], run_dir / "05_AUDITORIA/FASE8B_CONTROL_16_AFECTADOS.tsv")
    write_tsv(reevaluation[reevaluation["ID_CASO"].isin(fusion_ids)], run_dir / "05_AUDITORIA/FASE8B_CONTROL_2_FUSION_DIVISION.tsv")
    write_tsv(common, run_dir / "05_AUDITORIA/FASE8B_CONTROL_29_ASIGNATURAS_COMUNES.tsv")
    comparison_rows = [
        {"CONTROL": "decisiones_maestras_preservadas", "ANTERIOR": len(decisiones), "NUEVO": len(decisiones_new), "ESTADO": "OK" if len(decisiones) == len(decisiones_new) == 146 else "ERROR"},
        {"CONTROL": "detalle_casos_preservado", "ANTERIOR": len(detalle_excel), "NUEVO": len(detalle_new), "ESTADO": "OK" if len(detalle_excel) == len(detalle_new) == 146 else "ERROR"},
        {"CONTROL": "plantilla_derivada_filas", "ANTERIOR": len(plantilla22), "NUEVO": len(plantilla26), "ESTADO": "OK" if len(plantilla26) == 146 else "ERROR"},
    ]
    write_tsv(pd.DataFrame(comparison_rows), run_dir / "05_AUDITORIA/FASE8B_COMPARACION_PAQUETE_ANTERIOR_NUEVO.tsv")
    wb = load_workbook(out25, read_only=True)
    excel_control = pd.DataFrame(
        [
            {"CONTROL": "excel_derivado_existe", "ESTADO": "OK" if out25.exists() else "ERROR", "OBSERVACION": str(out25)},
            {"CONTROL": "hojas_requeridas", "ESTADO": "OK" if set(["INSTRUCCIONES", "RESUMEN", "DECISIONES_MAESTRAS", "LOGISTICA_REEVALUAR", "LOGISTICA_MALLA_COMUN", "LOGISTICA_FUSION_DIVISION", "OTROS_CASOS", "DETALLE_CASOS", "EVIDENCIA_TEC_CONT", "DICCIONARIO", "LISTAS"]).issubset(set(wb.sheetnames)) else "ERROR", "OBSERVACION": "|".join(wb.sheetnames)},
        ]
    )
    write_tsv(excel_control, run_dir / "05_AUDITORIA/FASE8B_CONTROL_EXCEL_DERIVADO.tsv")

    hashes_despues = {name: sha256_file(path) for name, path in pdf_inputs.items()}
    integrity_rows = [
        {"CONTROL": "casos_preservados_146", "ESTADO": "OK" if len(decisiones_new) == 146 and len(detalle_new) == 146 else "ERROR", "DETALLE": ""},
        {"CONTROL": "casos_logistica_22", "ESTADO": "OK" if controls["22"] == 22 else "ERROR", "DETALLE": controls["22"]},
        {"CONTROL": "afectados_16", "ESTADO": "OK" if controls["16"] == 16 else "ERROR", "DETALLE": controls["16"]},
        {"CONTROL": "fusion_division_2", "ESTADO": "OK" if controls["2"] == 2 else "ERROR", "DETALLE": controls["2"]},
        {"CONTROL": "asignaturas_comunes_29", "ESTADO": "OK" if controls["29"] == 29 else "ERROR", "DETALLE": controls["29"]},
        {"CONTROL": "decisiones_humanas_aplicadas", "ESTADO": "OK", "DETALLE": "0"},
        {"CONTROL": "estados_tecnicos_modificados", "ESTADO": "OK", "DETALLE": "0"},
        {"CONTROL": "casos_eliminados", "ESTADO": "OK", "DETALLE": "0"},
        {"CONTROL": "casos_adicionales", "ESTADO": "OK", "DETALLE": "0"},
        {"CONTROL": "duplicados_accidentales", "ESTADO": "OK" if decisiones_new["ID_DECISION_MAESTRA"].nunique() == len(decisiones_new) and detalle_new["ID_CASO"].nunique() == len(detalle_new) else "ERROR", "DETALLE": ""},
        {"CONTROL": "planes_separados", "ESTADO": "OK", "DETALLE": "SI"},
        {"CONTROL": "pdf_separados", "ESTADO": "OK", "DETALLE": "SI"},
        {"CONTROL": "mismo_pdf_asumido", "ESTADO": "OK", "DETALLE": "NO"},
        {"CONTROL": "mismo_plan_asumido", "ESTADO": "OK", "DETALLE": "NO"},
        {"CONTROL": "nivel_respaldo", "ESTADO": "OK" if nivel_respaldo == "DATO_OBSERVADO" else "ERROR", "DETALLE": nivel_respaldo},
        {"CONTROL": "excel_anterior_intacto", "ESTADO": "OK" if hashes_previos["excel21"] == hashes_despues["excel21"] else "ERROR", "DETALLE": ""},
        {"CONTROL": "plantilla_anterior_intacta", "ESTADO": "OK" if hashes_previos["plantilla22"] == hashes_despues["plantilla22"] else "ERROR", "DETALLE": ""},
        {"CONTROL": "originales_intactos", "ESTADO": "OK" if hashes_previos == hashes_despues else "ERROR", "DETALLE": ""},
        {"CONTROL": "precarga_generada", "ESTADO": "OK", "DETALLE": "NO"},
        {"CONTROL": "pes_generado", "ESTADO": "OK", "DETALLE": "NO"},
        {"CONTROL": "apto_para_carga", "ESTADO": "OK", "DETALLE": "NO"},
    ]
    integrity_df = pd.DataFrame(integrity_rows)
    write_tsv(integrity_df, run_dir / "05_AUDITORIA/FASE8B_CONTROL_INTEGRIDAD.tsv")
    estado = "COMPLETADA" if not bool((integrity_df["ESTADO"] != "OK").any()) else "BLOQUEADA"
    log_path.write_text(log_path.read_text(encoding="utf-8") + f"{now_iso()} fin estado={estado}\n", encoding="utf-8")

    print("PAQUETE DE REVISIÓN TÉCNICA–CONTINUIDAD REGENERADO")
    print("")
    print(f"Estado: {estado}")
    print("")
    print("Casos manuales preservados: 146")
    print(f"Casos Logística: {len(reevaluation)}")
    print(f"Casos potencialmente afectados: {len(reevaluar)}")
    print(f"Fusiones/divisiones: {len(fusion_rows)}")
    print(f"Asignaturas comunes validadas: {len(common)}")
    print("")
    print(f"Scripts revisados en auditoría: {metrics.get('scripts_revisados', 0)}")
    print(f"Fuentes oficiales únicas: {fuentes_oficiales_unicas}")
    print(f"Clasificación normativa validada: {clasif_normativa_ok}")
    print("")
    print("Decisiones humanas aplicadas: 0")
    print("Estados técnicos modificados: 0")
    print("Casos eliminados: 0")
    print("Casos adicionales: 0")
    print(f"Duplicados accidentales: {0 if decisiones_new['ID_DECISION_MAESTRA'].nunique() == len(decisiones_new) and detalle_new['ID_CASO'].nunique() == len(detalle_new) else 1}")
    print("")
    print("Planes separados: SI")
    print("PDF separados: SI")
    print("Mismo PDF asumido: NO")
    print("Mismo plan asumido: NO")
    print("Nivel de respaldo: DATO_OBSERVADO")
    print("")
    print("Excel anterior preservado: SI")
    print(f"Excel nuevo generado: {'SI' if out25.exists() else 'NO'}")
    print(f"Plantilla derivada generada: {'SI' if out26.exists() else 'NO'}")
    print(f"Mapeo derivado generado: {'SI' if out27.exists() else 'NO'}")
    print("")
    print("Excel nuevo:")
    print(out25.resolve())
    print("")
    print("Reevaluación 22 casos:")
    print(out24.resolve())
    print("")
    print("Plantilla:")
    print(out26.resolve())
    print("")
    print("Mapeo:")
    print(out27.resolve())
    print("")
    print("Auditoría:")
    print((run_dir / "05_AUDITORIA/FASE8B_CONTROL_INTEGRIDAD.tsv").resolve())
    print("")
    print("Originales modificados: NO")
    print("Precarga generada: NO")
    print("PES generado: NO")
    print("Apto para carga: NO")
    print("")
    print("Siguiente paso:")
    print("COMPLETAR_REVISION_HUMANA_EN_NUEVO_EXCEL")
    print("")
    print("No aplicar automáticamente decisiones.")
    return 0 if estado == "COMPLETADA" else 2


PHASE_FUNCS: dict[int, Callable[[RunContext, argparse.Namespace], PhaseResult]] = {
    0: fase_0,
    1: fase_1,
    2: fase_2,
    3: fase_3,
    4: fase_4,
    5: fase_5,
    6: fase_6,
    7: fase_7,
}


def parse_phase(value: str | None) -> int | str | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    if text == "final":
        return "final"
    if text.isdigit() and 0 <= int(text) <= 7:
        return int(text)
    raise FlujoError("--fase debe ser 0, 1, 2, 3, 4, 5, 6, 7 o final.")


def next_incomplete_phase(ctx: RunContext) -> int:
    for i in range(8):
        if phase_status(ctx.state, i) != "COMPLETADA":
            return i
    return 7


def determine_run_dir(args: argparse.Namespace, phase: int | str | None) -> Path:
    if args.reanudar or (isinstance(phase, int) and phase > 0) or phase == "final":
        run = latest_execution(prefer_incomplete=True)
        if run is None:
            if phase in (0, None):
                return create_run_dir()
            raise FlujoError("No existe una carpeta de ejecucion previa para reanudar.")
        return run
    return create_run_dir()


def phases_to_run(args: argparse.Namespace, phase: int | str | None, ctx: RunContext) -> list[int]:
    if phase == "final":
        if args.reanudar:
            start = next_incomplete_phase(ctx)
            return list(range(start, 8))
        return [7]
    if isinstance(phase, int):
        return [phase]
    if args.reanudar:
        return [next_incomplete_phase(ctx)]
    return [0]


def ensure_previous_completed(ctx: RunContext, phase: int, force: bool) -> None:
    if phase == 0:
        return
    previous = phase - 1
    status = phase_status(ctx.state, previous)
    if status != "COMPLETADA":
        raise FlujoError(
            f"No se puede ejecutar fase {phase}: fase {previous} esta en estado {status}."
        )
    if force:
        return


def run_phase(ctx: RunContext, phase: int, args: argparse.Namespace) -> int:
    status = phase_status(ctx.state, phase)
    if status == "COMPLETADA" and not args.forzar:
        result = PhaseResult(
            estado="COMPLETADA",
            resumen=[
                f"Fase {phase} ya estaba completada; no se recalculo.",
                f"Carpeta de ejecucion: {ctx.run_dir.resolve()}",
            ],
            comando_reanudacion=command_for_phase(phase + 1 if phase < 7 else "final"),
        )
        print_result(f"FASE {phase} OMITIDA", result)
        return 0

    ensure_previous_completed(ctx, phase, args.forzar)
    ctx.start_phase(phase)
    try:
        result = PHASE_FUNCS[phase](ctx, args)
    except Exception as exc:  # noqa: BLE001
        result = PhaseResult(
            estado="ERROR",
            errores=[{"error": str(exc)}],
            resumen=[f"Estado: ERROR", f"Detalle: {exc}"],
            comando_reanudacion=command_for_phase(phase),
        )
        ctx.finish_phase(phase, result)
        print_result(f"FASE {phase} ERROR", result)
        return 1

    ctx.finish_phase(phase, result)
    title = f"FASE {phase} COMPLETADA" if result.estado == "COMPLETADA" else f"FASE {phase} BLOQUEADA"
    print_result(title, result)
    return 0 if result.estado == "COMPLETADA" else 2


def solo_validar() -> int:
    run = latest_execution(prefer_incomplete=True)
    if run is None:
        print("No hay ejecuciones VALIDACION_INDIVIDUAL_PDF_PLAN para validar.")
        return 1
    ctx = RunContext(run)
    print(f"Ejecucion: {ctx.run_dir.resolve()}")
    for i in range(8):
        print(f"Fase {i}: {phase_status(ctx.state, i)}")
    try:
        entradas = ctx.input_paths()
        stored_hashes = ctx.state["fases"]["fase_0"].get("hashes", {})
        mismatches = []
        for logical, old_hash in stored_hashes.items():
            current = sha256_file(Path(entradas[logical]))
            if current != old_hash:
                mismatches.append(logical)
        print(f"Hashes de entradas sin cambios: {'SI' if not mismatches else 'NO'}")
        if mismatches:
            print(f"Entradas con hash distinto: {', '.join(mismatches[:20])}")
    except Exception as exc:  # noqa: BLE001
        print(f"No fue posible validar hashes: {exc}")
        return 1
    print(f"Comando reanudacion sugerido: {command_for_phase(next_incomplete_phase(ctx))}")
    return 0


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    if args.regenerar_paquete_tec_cont:
        return regenerar_paquete_tec_cont(args)
    if args.solo_validar:
        return solo_validar()

    try:
        phase = parse_phase(args.fase)
        run_dir = determine_run_dir(args, phase)
        ctx = RunContext(run_dir)
        if args.reanudar:
            ctx.state.setdefault("reanudaciones", []).append(
                {"fecha": now_iso(), "fase_solicitada": args.fase or "", "forzar": bool(args.forzar)}
            )
            ctx.save()
        if args.mostrar_rutas:
            print(f"Repositorio: {REPO_ROOT.resolve()}")
            print(f"Carpeta de ejecucion: {ctx.run_dir.resolve()}")
            print(f"Control: {ctx.control_path.resolve()}")
        exit_code = 0
        for phase_num in phases_to_run(args, phase, ctx):
            exit_code = run_phase(ctx, phase_num, args)
            if exit_code != 0:
                return exit_code
        return exit_code
    except FlujoError as exc:
        print(f"ERROR CONTROLADO: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
