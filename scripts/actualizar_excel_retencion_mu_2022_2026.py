from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
DESKTOP_XLSX = Path("/Users/alexi/Desktop/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026.xlsx")
SIES_XLSX = Path("/Users/alexi/Downloads/Informe_Retencion_SIES_2026_270702026.xlsx")
SIES_PDF = Path("/Users/alexi/Downloads/Informe-de-Retención-de-1er-Año-de-Pregrado-2026.pdf")
RUN_ID = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_DIR = REPO / "outputs" / "base_retencion_matricula_unificada_2022_2026" / RUN_ID

BASE_SHEET = "BASE_RETENCION_MU"
CALC_SHEET = "02_CALCULO_RETENCION"
EXEC_SHEET = "03_RESUMEN_EJECUTIVO"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def get_sies_ipss_general() -> dict[int, float]:
    wb = load_workbook(SIES_XLSX, read_only=True, data_only=True)
    ws = wb["Retención 1er año x IES"]
    values = {}
    header = None
    current_section = None
    for row in ws.iter_rows(values_only=True):
        first = row[0] if row else None
        if isinstance(first, str) and "Evolución de Retención de 1er año - IP" in first:
            current_section = "IP_GENERAL"
        elif isinstance(first, str) and "carreras técnicas - IP" in first:
            current_section = "IP_TECNICAS"
        elif isinstance(first, str) and "carreras profesionales - IP" in first:
            current_section = "IP_PROFESIONALES"
        if first == "Cod. Institución":
            header = list(row)
            continue
        if current_section == "IP_GENERAL" and first == 162:
            for idx, year in enumerate(header):
                if isinstance(year, int) and year >= 2022:
                    val = row[idx]
                    if isinstance(val, (int, float)):
                        values[year] = float(val)
            break
    wb.close()
    return values


def style_range(ws, cell_range: str, fill: str | None = None, bold: bool = False) -> None:
    thin = Side(style="thin", color="D9E2EC")
    for row in ws[cell_range]:
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if bold:
                cell.font = Font(bold=True)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        local = Path(tmp) / DESKTOP_XLSX.name
        shutil.copy2(DESKTOP_XLSX, local)
        wb = load_workbook(local)

        if CALC_SHEET in wb.sheetnames:
            del wb[CALC_SHEET]
        if EXEC_SHEET in wb.sheetnames:
            del wb[EXEC_SHEET]

        base = wb[BASE_SHEET]
        headers = [cell.value for cell in base[1]]
        col = {name: idx + 1 for idx, name in enumerate(headers)}
        max_row = base.max_row

        # Keep the existing first sheet data unchanged; only reinforce alignment/number formats.
        for row in base.iter_rows(min_row=1, max_row=max_row, max_col=base.max_column):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        calc = wb.create_sheet(CALC_SHEET, 1)
        calc_headers = [
            "ANIO_COHORTE",
            "ANIO_SIGUIENTE",
            "TIPO_DOC",
            "N_DOC",
            "DV",
            "SEXO",
            "COD_CAR",
            "VIG_ORIGEN",
            "ES_COHORTE_1ER_ANIO",
            "APARECE_VIGENTE_ANIO_SIGUIENTE",
            "RETENIDO_1ER_ANIO",
            "LLAVE_PERSONA",
        ]
        calc.append(calc_headers)
        base_last = max_row
        for r in range(2, max_row + 1):
            excel_row = [
                f"='{BASE_SHEET}'!{get_column_letter(col['ANIO_INFORMADO'])}{r}",
                f"=A{r}+1",
                f"='{BASE_SHEET}'!{get_column_letter(col['TIPO_DOC'])}{r}",
                f"='{BASE_SHEET}'!{get_column_letter(col['N_DOC'])}{r}",
                f"='{BASE_SHEET}'!{get_column_letter(col['DV'])}{r}",
                f"='{BASE_SHEET}'!{get_column_letter(col['SEXO'])}{r}",
                f"='{BASE_SHEET}'!{get_column_letter(col['COD_CAR'])}{r}",
                f"='{BASE_SHEET}'!{get_column_letter(col['VIG'])}{r}",
                f"=--(AND(A{r}='{BASE_SHEET}'!{get_column_letter(col['ANIO_ING_ACT'])}{r},'{BASE_SHEET}'!{get_column_letter(col['SEM_ING_ACT'])}{r}=1,H{r}=1,A{r}<MAX('{BASE_SHEET}'!$A$2:$A${base_last})))",
                f"=COUNTIFS('{BASE_SHEET}'!$A$2:$A${base_last},B{r},'{BASE_SHEET}'!${get_column_letter(col['TIPO_DOC'])}$2:${get_column_letter(col['TIPO_DOC'])}${base_last},C{r},'{BASE_SHEET}'!${get_column_letter(col['N_DOC'])}$2:${get_column_letter(col['N_DOC'])}${base_last},D{r},'{BASE_SHEET}'!${get_column_letter(col['DV'])}$2:${get_column_letter(col['DV'])}${base_last},E{r},'{BASE_SHEET}'!${get_column_letter(col['VIG'])}$2:${get_column_letter(col['VIG'])}${base_last},1)",
                f"=--(AND(I{r}=1,J{r}>0))",
                f"=C{r}&\"|\"&D{r}&\"|\"&E{r}",
            ]
            calc.append(excel_row)

        calc.freeze_panes = "A2"
        calc.auto_filter.ref = calc.dimensions
        style_range(calc, f"A1:L{max_row}", fill=None)
        style_range(calc, "A1:L1", fill="D9EAF7", bold=True)
        for idx in range(1, 13):
            calc.column_dimensions[get_column_letter(idx)].width = 18
        for idx in [1, 2, 4, 7, 8, 9, 10, 11]:
            for cell in calc.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=max_row):
                for c in cell:
                    c.number_format = "0"

        exec_ws = wb.create_sheet(EXEC_SHEET, 2)
        exec_ws.sheet_view.showGridLines = False
        exec_ws["A1"] = "Resumen ejecutivo - Retención de 1er año IPSS"
        exec_ws["A2"] = "Institución 162 - cálculo interno desde Matrícula Unificada informada a SIES y comparación con informe SIES 2026"
        exec_ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
        exec_ws["A2"].font = Font(color="FFFFFF")
        for row in exec_ws["A1:H2"]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")
        exec_ws.merge_cells("A1:H1")
        exec_ws.merge_cells("A2:H2")

        table_headers = ["Cohorte", "Matricula 1er año", "Retenidos año siguiente", "Retención interna", "Retención SIES IPSS", "Brecha pp", "Mujeres", "Hombres"]
        for c, h in enumerate(table_headers, 1):
            exec_ws.cell(5, c, h)
        years = [2022, 2023, 2024, 2025]
        sies = get_sies_ipss_general()
        for i, y in enumerate(years, 6):
            exec_ws.cell(i, 1, y)
            exec_ws.cell(i, 2, f"=SUMIFS('{CALC_SHEET}'!$I$2:$I${max_row},'{CALC_SHEET}'!$A$2:$A${max_row},A{i})")
            exec_ws.cell(i, 3, f"=SUMIFS('{CALC_SHEET}'!$K$2:$K${max_row},'{CALC_SHEET}'!$A$2:$A${max_row},A{i})")
            exec_ws.cell(i, 4, f"=IFERROR(C{i}/B{i},\"\")")
            exec_ws.cell(i, 5, sies.get(y))
            exec_ws.cell(i, 6, f"=IFERROR((D{i}-E{i})*100,\"\")")
            exec_ws.cell(i, 7, f"=SUMIFS('{CALC_SHEET}'!$I$2:$I${max_row},'{CALC_SHEET}'!$A$2:$A${max_row},A{i},'{CALC_SHEET}'!$F$2:$F${max_row},\"M\")")
            exec_ws.cell(i, 8, f"=SUMIFS('{CALC_SHEET}'!$I$2:$I${max_row},'{CALC_SHEET}'!$A$2:$A${max_row},A{i},'{CALC_SHEET}'!$F$2:$F${max_row},\"H\")")

        exec_ws["A12"] = "Lectura ejecutiva"
        exec_ws["A13"] = "La tasa interna se calcula con fórmulas sobre la hoja base: cohorte vigente de primer año del año Y y aparición vigente de la misma persona en Y+1."
        exec_ws["A14"] = "La comparación SIES usa la fila IP SAN SEBASTIAN, código 162, sección Evolución de Retención de 1er año - IP del Excel oficial adjunto."
        exec_ws["A15"] = "Los datos de 2022 y 2023 no contienen FECHA_MATRICULA ni REINCORPORACION equivalentes al esquema 2026; esa limitación queda en gobernanza."
        exec_ws.merge_cells("A12:H12")
        exec_ws.merge_cells("A13:H13")
        exec_ws.merge_cells("A14:H14")
        exec_ws.merge_cells("A15:H15")
        style_range(exec_ws, "A5:H10", fill=None)
        style_range(exec_ws, "A5:H5", fill="D9EAF7", bold=True)
        style_range(exec_ws, "A12:H15", fill="F4F7FA")
        exec_ws["A12"].font = Font(bold=True)
        for r in range(6, 10):
            for c in [4, 5]:
                exec_ws.cell(r, c).number_format = "0.0%"
            exec_ws.cell(r, 6).number_format = "0.0"
        for c in range(1, 9):
            exec_ws.column_dimensions[get_column_letter(c)].width = 19

        line = LineChart()
        line.title = "Retención interna vs SIES"
        line.y_axis.title = "Tasa"
        line.x_axis.title = "Cohorte"
        line.add_data(Reference(exec_ws, min_col=4, max_col=5, min_row=5, max_row=9), titles_from_data=True)
        line.set_categories(Reference(exec_ws, min_col=1, min_row=6, max_row=9))
        line.height = 8
        line.width = 15
        exec_ws.add_chart(line, "A18")

        bar = BarChart()
        bar.title = "Cohorte y retenidos"
        bar.y_axis.title = "Registros"
        bar.x_axis.title = "Cohorte"
        bar.add_data(Reference(exec_ws, min_col=2, max_col=3, min_row=5, max_row=9), titles_from_data=True)
        bar.set_categories(Reference(exec_ws, min_col=1, min_row=6, max_row=9))
        bar.height = 8
        bar.width = 15
        exec_ws.add_chart(bar, "I18")

        for ws in [base, calc, exec_ws]:
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=cell.alignment.wrap_text)

        out_xlsx = OUT_DIR / f"BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_CON_FORMULAS_{RUN_ID}.xlsx"
        wb.save(out_xlsx)
        wb.close()

    shutil.copy2(out_xlsx, DESKTOP_XLSX)
    manifest = {
        "estado": "BASE_RETENCION_MU_2022_2026_FORMULAS_RETENCION_AGREGADAS",
        "timestamp": RUN_ID,
        "proceso": "Matricula Unificada SIES",
        "subproyecto": "Calculo retencion primer anio respecto al anio anterior informado",
        "fuentes_referencia": [
            {"archivo": str(SIES_PDF), "clasificacion": "informe SIES PDF de referencia visual/conceptual", "sha256": sha256(SIES_PDF)},
            {"archivo": str(SIES_XLSX), "clasificacion": "Excel SIES con tablas de comparacion", "sha256": sha256(SIES_XLSX)},
        ],
        "salidas": {
            "excel_repositorio": str(out_xlsx),
            "excel_escritorio": str(DESKTOP_XLSX),
            "sha256_excel_repositorio": sha256(out_xlsx),
            "sha256_excel_escritorio": sha256(DESKTOP_XLSX),
        },
        "definicion_calculo": {
            "cohorte": "VIG=1, ANIO_INFORMADO=Y, ANIO_ING_ACT=Y, SEM_ING_ACT=1",
            "retenido": "Misma persona TIPO_DOC+N_DOC+DV aparece con VIG=1 en ANIO_INFORMADO=Y+1",
            "tasa": "retenidos / cohorte",
            "nota": "Calculo tecnico interno sobre Matricula Unificada informada; comparacion SIES usa informe adjunto.",
        },
        "comparacion_sies_ipss_162": get_sies_ipss_general(),
    }
    (OUT_DIR / "MANIFEST_RETENCION_CON_FORMULAS.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    report = [
        "REPORTE VALIDACION RETENCION CON FORMULAS",
        f"Timestamp: {RUN_ID}",
        f"Excel repositorio: {out_xlsx}",
        f"Excel escritorio: {DESKTOP_XLSX}",
        f"SHA-256 repositorio: {manifest['salidas']['sha256_excel_repositorio']}",
        f"SHA-256 escritorio: {manifest['salidas']['sha256_excel_escritorio']}",
        "Hojas: BASE_RETENCION_MU, 02_CALCULO_RETENCION, 03_RESUMEN_EJECUTIVO",
        "La primera hoja conserva la base ya construida; se agregan calculo y resumen con formulas.",
        "No se modificaron fuentes originales.",
    ]
    (OUT_DIR / "REPORTE_VALIDACION_RETENCION_CON_FORMULAS.txt").write_text("\n".join(report), encoding="utf-8")
    with zipfile.ZipFile(out_xlsx) as zf:
        bad = zf.testzip()
    print(json.dumps({"out_dir": str(OUT_DIR), "xlsx": str(out_xlsx), "desktop": str(DESKTOP_XLSX), "testzip": bad, "sha": sha256(out_xlsx)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
