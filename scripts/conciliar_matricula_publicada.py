#!/usr/bin/env python3
"""Recuenta fuentes MU2026 y la publicacion adjunta, sin ejecutar el motor."""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from common.catalogos_sies import escribir_inmutables, json_bytes, normalizar, sha256

# Esquemas MU2026, documentados en comparar_publicacion2026_integral_pregrado_posgrado.py.
BASE = "TIPO_DOC N_DOC DV PRIMER_APELLIDO SEGUNDO_APELLIDO NOMBRE SEXO FECH_NAC NAC PAIS_EST_SEC COD_SED COD_CAR MODALIDAD JOR VERSION FOR_ING_ACT ANIO_ING_ACT SEM_ING_ACT ANIO_ING_ORI SEM_ING_ORI".split()
PRE_COLUMNS = BASE + "ASI_INS_ANT ASI_APR_ANT PROM_PRI_SEM PROM_SEG_SEM ASI_INS_HIS ASI_APR_HIS NIV_ACA SIT_FON_SOL SUS_PRE FECHA_MATRICULA REINCORPORACION VIG".split()
POST_COLUMNS = BASE + ["VIG"]


def entero(value: object) -> int:
    if value is None or isinstance(value, bool):
        raise ValueError("Conteo/codigo ausente o booleano")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError("Conteo/codigo no numerico") from exc
    if not number.is_finite() or number != number.to_integral_value() or number < 0:
        raise ValueError("Conteo/codigo no entero no negativo")
    return int(number)


def resumir(rows: list[dict]) -> dict:
    if not rows:
        raise ValueError("Fuente institucional vacia")
    vig, personas, modalidades, sedes = Counter(), set(), Counter(), Counter()
    for numero, row in enumerate(rows, 1):
        state = entero(row["VIG"])
        if state not in {0, 1, 2}:
            raise ValueError(f"VIG fuera de dominio, registro {numero}")
        vig[str(state)] += 1
        if state in {1, 2}:
            identity = tuple(re.sub(r"^(\d+)\.0+$", r"\1", normalizar(row[k]))
                             for k in ("TIPO_DOC", "N_DOC", "DV"))
            if any(v in {"", "NAN", "NONE", "NULL"} for v in identity):
                raise ValueError(f"Identidad incompleta, registro {numero}")
            personas.add(identity)
            modalidad = entero(row["MODALIDAD"])
            if modalidad not in {1, 2, 3}:
                raise ValueError(f"Modalidad fuera de dominio, registro {numero}")
            modalidades[str(modalidad)] += 1
            sedes[str(entero(row["COD_SED"]))] += 1
    activos = vig["1"] + vig["2"]
    return {"registros": len(rows), "vig": {k: vig[k] for k in ("0", "1", "2")},
            "matriculas_comparables": activos, "personas_comparables": len(personas),
            "matriculas_adicionales_respecto_personas": activos - len(personas),
            "matriculas_por_modalidad": dict(sorted(modalidades.items())),
            "matriculas_por_sede": dict(sorted(sedes.items()))}


def leer_pregrado(path: Path) -> list[dict]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book["CONSOLIDADO"]
        if sheet.max_column != len(PRE_COLUMNS):
            raise ValueError("Pregrado: se esperan 32 columnas")
        rows = sheet.iter_rows(values_only=True)
        if list(next(rows)) != PRE_COLUMNS:
            raise ValueError("Pregrado: encabezados distintos del contrato MU2026")
        return [dict(zip(PRE_COLUMNS, row)) for row in rows if any(v is not None for v in row)]
    finally:
        book.close()


def leer_postitulo(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream, delimiter=";"))
    if any(len(row) != len(POST_COLUMNS) for row in rows):
        raise ValueError("Postitulo: se esperan 21 campos sin encabezado")
    if rows and rows[0][0] == "TIPO_DOC":
        raise ValueError("Postitulo: encabezado inesperado")
    return [dict(zip(POST_COLUMNS, row)) for row in rows]


def leer_publicacion(path: Path, year: int) -> dict:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        name = f"Listado de instituciones {year}"
        sheet = book[name]
        expected = ["Nombre de la institucion", "Codigo institucional",
                    f"Matricula Total Pregrado {year}", f"Matricula Total Posgrado {year}",
                    f"Matricula Total Postitulo {year}", f"Matricula Total {year}"]
        header = next(sheet.iter_rows(min_row=3, max_row=3, max_col=6, values_only=True))
        if list(map(normalizar, header)) != list(map(normalizar, expected)):
            raise ValueError("Publicacion: encabezados no reconocidos")
        matches = [(i, row) for i, row in enumerate(sheet.iter_rows(min_row=4, max_col=6, values_only=True), 4)
                   if normalizar(row[1]) in {"162", "162.0"}]
        if len(matches) != 1:
            raise ValueError("Publicacion: IPSS ausente o repetido")
        index, row = matches[0]
        if normalizar(row[0]) != "IP SAN SEBASTIAN":
            raise ValueError("Publicacion: nombre y codigo institucional inconsistentes")
        pre, post, total = entero(row[2]), entero(row[4]), entero(row[5])
        posgrado = None if row[3] is None else entero(row[3])
        if pre + post + (posgrado or 0) != total:
            raise ValueError("Publicacion: total no concilia con los componentes disponibles")
        return {"pregrado": pre, "posgrado": posgrado, "postitulo": post, "total": total,
                "celda_posgrado_vacia": posgrado is None, "hoja": name,
                "localizador": f"B{index}:F{index}"}
    finally:
        book.close()


def conciliar(pregrado: Path, postitulo: Path, publicacion: Path, *, anio: int, fecha: str) -> dict:
    if anio != 2026:
        raise ValueError("Solo existe contrato de conciliacion MU2026; no reutilizarlo para otro ano")
    date.fromisoformat(fecha)
    paths = {"pregrado": pregrado, "postitulo": postitulo, "publicacion": publicacion}
    hashes = {k: sha256(p) for k, p in paths.items()}
    pre, post = resumir(leer_pregrado(pregrado)), resumir(leer_postitulo(postitulo))
    pub = leer_publicacion(publicacion, anio)
    if any(sha256(p) != hashes[k] for k, p in paths.items()):
        raise ValueError("Una fuente cambio durante la lectura")
    diferencias = {"pregrado": pub["pregrado"] - pre["matriculas_comparables"],
                   "postitulo": pub["postitulo"] - post["matriculas_comparables"],
                   "total": pub["total"] - pre["matriculas_comparables"] - post["matriculas_comparables"]}
    return {
        "schema_version": 1, "anio": anio, "fecha_verificacion": fecha,
        "estado": "DISCREPANCIA_NO_EXPLICADA" if any(diferencias.values()) else "COINCIDENCIA_AGREGADA_NO_CERTIFICA_CARGA",
        "fuentes": {k: {"archivo": p.name, "sha256": hashes[k]} for k, p in paths.items()},
        "pregrado_institucional": pre, "postitulo_institucional": post,
        "publicacion": pub, "diferencia_publicacion_menos_institucional": diferencias,
        "regla": "Comparar matriculas VIG=1/2, no personas; VIG=0 solo control historico",
        "alcance": "Consolidado institucional reconstruido y copia aportada de publicacion; no certifica cadena de cargas aceptadas",
        "limitaciones": ["La diferencia agregada no identifica personas faltantes ni demuestra una causa",
                         "Posgrado vacio se conserva como null, no se publica como cero observado",
                         "El glosario adjunto OFICIAL_GLOSARIO_MATRICULA_WEB_E.pdf corresponde a 2025; no es manual 2026",
                         "No demuestra cobertura del puente CODCARPR ni habilita MU2027"],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for field in ("pregrado", "postitulo", "publicacion", "salida"):
        parser.add_argument(f"--{field}", type=Path, required=True)
    parser.add_argument("--anio", type=int, choices=[2026], required=True)
    parser.add_argument("--fecha", required=True)
    args = vars(parser.parse_args())
    output = args.pop("salida")
    report = conciliar(**args)
    escribir_inmutables({output: json_bytes(report)})
    print(json.dumps({"estado": report["estado"], "diferencias": report["diferencia_publicacion_menos_institucional"]}))


if __name__ == "__main__":
    main()
