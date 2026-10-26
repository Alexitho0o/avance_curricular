"""Lee respaldo T3 recuperado y contrasta F02 sin adoptar versiones ni cambiar bases."""
import json,csv,io,hashlib,datetime,collections,subprocess,sys
from pathlib import Path
from zoneinfo import ZoneInfo
import openpyxl
C=json.load(open(sys.argv[1]));O=Path(C['output']);P=Path(C['privado']);A=Path(C['anterior']);AP=Path(C['privado_anterior']);R=O.parents[1]
load=lambda p:json.load(open(p))
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str))
def val(v):return '' if v is None else str(int(v)) if isinstance(v,float) and v.is_integer() else str(v).strip()
def norm(v):return val(v).upper().replace('.','').replace(' ','')
def date(v):
 if isinstance(v,(datetime.datetime,datetime.date)):return v.strftime('%Y-%m-%d')
 s=val(v)
 for f in ['%Y-%m-%d %H:%M:%S','%Y-%m-%d','%d-%m-%Y','%d-%m-%y','%d/%m/%Y']:
  try:return datetime.datetime.strptime(s,f).strftime('%Y-%m-%d')
  except ValueError:pass
 return None
m=load(A/'MANIFEST.json');sourcehash={}
for name in ['REVISION_16_CASOS_FINAL.json','PERFIL_FUENTES_ADICIONALES.json']:
 assert sha(AP/name)==m['salidas_privadas_hashes'][name];sourcehash[str(AP/name)]=sha(AP/name)
rev=load(AP/'REVISION_16_CASOS_FINAL.json');f=next(x for x in rev if x['caso']=='F02');foco={x['persona']:x['caso'] for x in rev};owners=collections.defaultdict(set)
for x in rev:
 for s in x['segmentos_previos']:owners[s['CODCLI']].add(x['persona'])
assert all(len(x)==1 for x in owners.values())
prof=load(AP/'PERFIL_FUENTES_ADICIONALES.json');pp=next(x for x in prof if Path(x['path']).name=='PRUEBA CUMPLIMIENTO INGRESO CALIFICACIONES T3_2024.xlsx');path=Path(pp['path']);b=path.read_bytes();h=hashlib.sha256(b).hexdigest();assert h==pp['sha256'];sourcehash[str(path)]=h
w=openpyxl.load_workbook(io.BytesIO(b),read_only=True,data_only=True);rows=[];inventory=[];rejected=[];summaries=[]
for ws in w:
 it=iter(ws.values)
 if ws.title=='RESUMEN':
  n=0
  for rn,vals in enumerate(it,1):
   if any(v is not None for v in vals):summaries.append({'fila':rn,'valores':[val(v) for v in vals]});n+=1
  inventory.append({'hoja':ws.title,'filas_no_vacias':n,'dimension':[ws.max_row,ws.max_column],'uso':'Inspección completa de resumen; no es llave individual'});continue
 heads=[val(v).upper() for v in next(it)];n=0;found=0
 for rn,vals in enumerate(it,2):
  if not any(v is not None for v in vals):continue
  n+=1;d=dict(zip(heads,vals));cc=val(d.get('CODCLI'));rut=norm(d.get('RUT'));candidates=owners.get(cc,set());match=[person for person in candidates if rut in [person,person.replace('-',''),person.split('-')[0]]]
  # Secondary exact RUT body search is audit-only if CODCLI not previously located.
  if not candidates:
   audit=[person for person in foco if rut in [person,person.replace('-',''),person.split('-')[0]]]
   if audit:rejected.append({'hoja':ws.title,'fila':rn,'CODCLI':cc,'RUT_original':val(d.get('RUT')),'motivo':'Persona exacta, expediente no presente en inventario previo','personas':audit})
   continue
  if len(match)!=1:rejected.append({'hoja':ws.title,'fila':rn,'CODCLI':cc,'motivo':'RUT no corrobora CODCLI'});continue
  person=match[0];keep=['ANO','PERIODO','CODCLI','CODIGO_PLAN_ESTUDIO','CODCARRERA_ALUMNO','CODIGO_ASIGNATURA','CODIGO_CARRERA_ASIGNATURA','ASIGNATURA','SECCION','NOTAS_NO_INGRESADAS','NOTAS_INGRESADAS','TOTAL_A_INGRESAR','FECHAINICIO','FECHAFIN','ESTADO','ESTADO_ACADEMICO','FECHA MEDICION']
  z={k:val(d.get(k)) for k in keep if k in d};z.update(caso=foco[person],persona=person,RUT_original=val(d['RUT']),archivo=str(path),sha256=h,hoja=ws.title,fila=rn,metodo='CODCLI exacto + cuerpo RUT corroborado; DIG docente excluido',fechas_normalizadas={k:date(d.get(k)) for k in ['FECHAINICIO','FECHAFIN','FECHA MEDICION'] if k in d});rows.append(z);found+=1
 inventory.append({'hoja':ws.title,'filas_datos':n,'encabezados':heads,'filas_foco':found,'personas_foco':len({r['persona'] for r in rows if r['hoja']==ws.title})})
w.close();assert sha(path)==h
# Compare T3 observation against already-read master, without calling copied rows independent evidence.
comparison=[]
for r in rows:
 prior=next(x for x in rev if x['caso']==r['caso'])
 matches=[]
 for a in prior['filas_recuperadas_codcli']:
  if Path(a['archivo']).name!='CUMPLIMIENTO INGRESO CALIFICACIONES.xlsx':continue
  keys=['CODCLI','CODIGO_PLAN_ESTUDIO','CODCARRERA_ALUMNO','ANO','PERIODO','CODIGO_ASIGNATURA','SECCION','FECHAINICIO','FECHAFIN','ESTADO','ESTADO_ACADEMICO']
  if all(a.get(k,'')==r.get(k,'') for k in keys) and date(a.get('FECHA REPORTE'))==r['fechas_normalizadas'].get('FECHA MEDICION'):
   matches.append({'archivo':a['archivo'],'sha256':a['sha256'],'hoja':a['hoja'],'fila':a['fila']})
 comparison.append({'respaldo':r,'maestro_coincidente':matches,'tipo':'REPRODUCE_OBSERVACION_MAESTRO' if matches else 'OBSERVACION_ADICIONAL'})
# Re-read the earlier reports for F02, all academic rows, no names or unrelated identifiers.
base=load(R/'outputs/trayectorias_374_20261023_072401/MANIFEST.json');paths={str(Path(a['archivo'])):a['sha256'] for a in f['filas_adicionales'] if a['rol'].startswith('PROMEDIOS')}
for p,hh in base['fuentes'].items():
 if Path(p).name in ['PROMEDIOSDEALUMNOS_12166.xlsx','PROMEDIOSDEALUMNOS_7804.xlsx']:paths[p]=hh
pr=[];prinventory=[]
for p,hh in paths.items():
 data=Path(p).read_bytes();actual=hashlib.sha256(data).hexdigest();assert actual==hh;sourcehash[p]=actual
 wb=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True);ws=wb['Hoja1'];it=iter(ws.values);heads=[val(v).upper() for v in next(it)];n=found=0
 for rn,vals in enumerate(it,2):
  n+=1;d=dict(zip(heads,vals));rut=norm(d.get('RUT'));dv=norm(d.get('DIG'));person=rut if '-' in rut else rut+'-'+dv
  if person!=f['persona']:continue
  keep=['CODCLI','RUT','DIG','CODCARR','CARRERA','JORNADA','REGIMEN','RAMOEQUIV','CODRAMO','CODSECC','ASIGNATURA','EXAMEN','NOTA_FINAL','ANO','PERIODO','ESTADO','DESCRIPCION_ESTADO','CONVALIDADO','NIVEL','ESTADO_ACADEMICO','PLAN_DE_ESTUDIO']
  z={k:val(d.get(k)) for k in keep if k in d};z.update(archivo=p,sha256=hh,hoja='Hoja1',fila=rn);pr.append(z);found+=1
 prinventory.append({'archivo':p,'sha256':hh,'hoja':'Hoja1','filas_datos':n,'filas_F02':found,'corte_extraccion':'No inferido de nombre ni carpeta'});wb.close()
 print(json.dumps({'fuente':Path(p).name,'foco':found}),flush=True)
# Plan/date evidence visible in report, preserving precise locators only privately.
t3f=[r for r in rows if r['caso']=='F02'];priorf=f['filas_recuperadas_codcli'];dates=[]
for a in priorf:
 if a['rol']=='CALIFICACIONES_ASIGNATURA' and a.get('CODIGO_PLAN_ESTUDIO'):
  dates.append({'plan':a['CODIGO_PLAN_ESTUDIO'],'ano':a['ANO'],'periodo':a['PERIODO'],'fecha_reporte':date(a.get('FECHA REPORTE')),'inicio_asignatura':date(a.get('FECHAINICIO')),'fin_asignatura':date(a.get('FECHAFIN')),'archivo':a['archivo'],'hoja':a['hoja'],'fila':a['fila'],'CODCLI':a['CODCLI']})
for a in t3f:dates.append({'plan':a['CODIGO_PLAN_ESTUDIO'],'ano':a['ANO'],'periodo':a['PERIODO'],'fecha_reporte':a['fechas_normalizadas'].get('FECHA MEDICION'),'inicio_asignatura':a['fechas_normalizadas'].get('FECHAINICIO'),'fin_asignatura':a['fechas_normalizadas'].get('FECHAFIN'),'archivo':a['archivo'],'hoja':a['hoja'],'fila':a['fila'],'CODCLI':a['CODCLI']})
# The old-plan row that generated ambiguity must be preserved even if absent from T3.
old=[r for r in pr if r.get('PLAN_DE_ESTUDIO')=='IECIREOL20211' and r.get('ANO')=='2024' and r.get('PERIODO')=='2']
dump(P/'F02_ANTECEDENTE.json',f);dump(P/'FILAS_T3_FOCO_16.json',rows);dump(P/'COMPARACION_T3_MAESTRO.json',comparison);dump(P/'CONTROL_IDENTIDAD.json',rejected);dump(P/'RESUMEN_HOJA_ORIGINAL.json',summaries);dump(P/'F02_PROMEDIOS_RELECTURA.json',pr);dump(P/'F02_CRONOLOGIA_FUENTES.json',dates);dump(P/'F02_PLAN_ANTERIOR_2024_2.json',old)
summary={'fuente_T3_hash':h,'tres_hojas_leidas':len(inventory)==3,'inventario_T3':inventory,'filas_foco_T3':len(rows),'personas_foco_T3':len({r['persona'] for r in rows}),'filas_F02_T3':len(t3f),'planes_F02_T3':dict(collections.Counter(r['CODIGO_PLAN_ESTUDIO'] for r in t3f)),'fechas_medicion_F02_T3':sorted({r['fechas_normalizadas'].get('FECHA MEDICION') for r in t3f if r['fechas_normalizadas'].get('FECHA MEDICION')}),'filas_F02_T3_reproducidas_maestro':sum(x['respaldo']['caso']=='F02' and bool(x['maestro_coincidente']) for x in comparison),'filas_F02_T3_adicionales':sum(x['respaldo']['caso']=='F02' and not x['maestro_coincidente'] for x in comparison),'filas_plan_anterior_2024_2':len(old),'fuentes_plan_anterior_2024_2':dict(collections.Counter(Path(r['archivo']).name for r in old)),'alertas_identidad':len(rejected),'fechas_reporte_plan_nuevo':sorted({r['fecha_reporte'] for r in dates if r['plan']=='IINF20241' and r['fecha_reporte']}),'fecha_efectiva_cambio':'Pendiente: ningún campo leído identifica un evento administrativo de cambio','codigo_anual_2024':f['antecedente']['codigos_anuales'],'codigos_nuevos_adoptados':0,'denominadores_conservados':[5901,6622,24774],'clasificacion':'B datos; C comparación; E fecha efectiva pendiente'}
dump(O/'RESULTADOS.json',summary);dump(O/'INVENTARIO_PROMEDIOS.json',prinventory);dump(O/'FUENTES_HASHES.json',sourcehash)
assert all(sha(p)==hh for p,hh in sourcehash.items());print(json.dumps(summary,ensure_ascii=False),flush=True)
