"""C: edición agregada T3-2025, preservación OOXML y QA; fuentes solo lectura."""
import argparse, json, hashlib, re, zipfile, shutil, subprocess
from pathlib import Path
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
import openpyxl
R=Path(__file__).resolve().parents[1]
O=R/'outputs/correccion_proxy_t3_2025_20261023_v15'
SRC=R/'outputs/correccion_proxy_t2_2025_20261023_v14/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V14.xlsx'
NAME='RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V15.xlsx'
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q=lambda s:'{'+NS+'}'+s
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(n,x): (O/n).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def git():return {k:subprocess.check_output(['git',*v],cwd=R,text=True).strip() for k,v in [('rama',['branch','--show-current']),('HEAD',['rev-parse','HEAD'])]}
def prepare():
    assert sha(SRC)=='20813b21f546085f94c10b2983335567c937e492f4d3f707be0893e5cb5234a4'
    r=json.loads((O/'RESULTADO_AGREGADO.json').read_text());assert r['totales']['numerador_propuesto']==546
    for p,v in r['fuentes'].items():assert sha(p)==v['sha256']
    w=openpyxl.load_workbook(SRC,read_only=True,data_only=False)
    edits=[];s='26_Avance_Curricular_Resumen'
    def add(sheet,cell,value=None,formula=None,expected=None,style_from=None):
        e=dict(sheet=sheet,cell=cell,before=w[sheet][cell].value)
        e.update(dict(value=value,expected=value) if formula is None else dict(formula=formula,expected=expected))
        if style_from:e['style_from']=style_from
        edits.append(e)
    add(s,'D1853',formula='=SUM(D1859:D1860)',expected=596)
    add(s,'E1853',formula='=SUM(E1859:E1860)',expected=546)
    add(s,'F1853',formula=w[s]['F1853'].value,expected=546/1150)
    add(s,'H1853','C · R/P incluidos; detalle debajo.')
    add(s,'A1857','T3-2025 · R/P',style_from='A1848')
    headers=['Grupo T1','Documento','Personas cohorte','Con registro T3','VIGENTE T3','Tasa en grupo','Fuente','Método']
    data=[['2025','R',1130,582,535,None,'PROMEDIOS / Hoja1','RUT + DIG'],['2025','P',20,14,11,None,'IPSS_Procesamiento','2025_carga: J / BU / BV'],['2025','Total',None,None,None,None,'ANO=2025; PERIODO=3','Una persona, un conteo']]
    for j,v in enumerate(headers,1):add(s,f'{openpyxl.utils.get_column_letter(j)}1858',v,style_from=f'{openpyxl.utils.get_column_letter(j)}1849')
    for i,row in enumerate(data,1859):
        for j,v in enumerate(row,1):
            col=openpyxl.utils.get_column_letter(j);a=f'{col}{i}';st=f'{col}1853'
            if j==6:add(s,a,formula=f'=IF(C{i}=0,"",E{i}/C{i})',expected=[535/1130,11/20,546/1150][i-1859],style_from=st)
            elif i==1861 and 3<=j<=5:add(s,a,formula=f'=SUM({col}1859:{col}1860)',expected={3:1150,4:596,5:546}[j],style_from=st)
            else:add(s,a,v,style_from=st)
    ss='52_Codigos_SIES_Revision'
    add(ss,'A1','52 · Seguimiento de T1 en T2 y T3 hasta 2026\nC/D · Se conserva el código T1, las bases y los denominadores. V15 completa los auxiliares PROMEDIOS de 2025 con documentos R/P. La revisión del código SIES permanece separada.')
    add(ss,'A14','C · PROMEDIOS 2025: T2=611/1.150 y T3=546/1.150 (hoja 26). D · Actividad del período + estado VIGENTE del extracto; no se exige fecha histórica para este análisis interno. Foto de matrícula T3: 1.023/1.150.')
    add(ss,'A32','C · V15 conserva diccionario v7, denominadores, tablas y gráficos. Corrige el auxiliar T3-2025 mediante el puente R/P y normaliza los estilos que Excel reparaba. La foto T3 no cambia.')
    add('18_Foto_Intertrimestral','B82','B · Foto T3: 2025 = Matrícula T3_2025 2.xlsx (hoja 24); 2026 = DATOSDEALUMNOS_12159 al 22-10-2026, T3 en curso (hoja 25). C/D · T2 del grupo 1 usa actividad del período y estado VIGENTE de PROMEDIOS como seguimiento interno.')
    add('28_QA_Matematica','A18','Alcance del cálculo interno')
    add('28_QA_Matematica','A19','C/D · Para el seguimiento interno se cuenta una persona con actividad del período y estado VIGENTE en PROMEDIOS, sin exigir fecha histórica del estado. Se deduplican las filas por asignatura. La foto de matrícula T3 mantiene su fuente y cálculo separados.')
    dump('EDICION.json',dict(source=str(SRC),source_sha256=sha(SRC),scratch='/tmp/retencion_proxy_t3_v15_20261023',edits=edits,git_antes=git(),fuentes=r['fuentes']))
    print('Ediciones preparadas:',len(edits))
def finish():
    m=json.loads((O/'EDICION.json').read_text());assert sha(SRC)==m['source_sha256']
    z=zipfile.ZipFile(SRC);a=zipfile.ZipFile(Path(m['scratch'])/'artifact_recalculado.xlsx')
    names=[s.get('name') for s in ET.fromstring(z.read('xl/workbook.xml')).find(Q('sheets'))]
    strings=[''.join(n.itertext()) for n in ET.fromstring(a.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in a.namelist() else []
    repl={};by={}
    for e in m['edits']:by.setdefault(e['sheet'],[]).append(e)
    for s,edits in by.items():
        part=f'xl/worksheets/sheet{names.index(s)+1}.xml';txt=z.read(part).decode()
        oldtree=ET.fromstring(txt);oldcells={c.get('r'):c for c in oldtree.iter(Q('c'))}
        cells={c.get('r'):c for c in ET.fromstring(a.read(part)).iter(Q('c'))};newrows={}
        for e in edits:
            address=e['cell'];c=cells[address];attrs={'r':address}
            ref=oldcells.get(address,oldcells.get(e.get('style_from')))
            if ref is not None:attrs.update(ref.attrib);attrs['r']=address
            attrs.pop('t',None)
            if c.get('t'):attrs['t']=c.get('t')
            if c.get('t')=='s':
                attrs['t']='inlineStr';body='<is><t xml:space="preserve">'+escape(strings[int(c.find(Q('v')).text)])+'</t></is>'
            else:
                body=''.join(ET.tostring(n,encoding='unicode').replace('ns0:','').replace(':ns0','') for n in c)
                body=re.sub(r' xmlns="[^"]+"','',body)
            new='<c '+' '.join(k+'="'+escape(v,{'"':'&quot;'})+'"' for k,v in attrs.items())+'>'+body+'</c>'
            pattern=rf'<c\b(?=[^>]*\br="{address}")[^>]*>.*?</c>'
            hit=re.search(pattern,txt,re.S)
            if hit:txt=txt[:hit.start()]+new+txt[hit.end():]
            else:newrows.setdefault(int(re.search(r'\d+',address).group()),[]).append(new)
        if newrows:
            assert s=='26_Avance_Curricular_Resumen'
            rows=''.join('<row r="'+str(r)+'">'+''.join(cs)+'</row>' for r,cs in sorted(newrows.items()))
            txt=txt.replace('</sheetData>',rows+'</sheetData>')
            txt=re.sub(r'<dimension ref="[^"]+"\s*/>', '<dimension ref="A1:H1861"/>',txt)
        repl[part]=txt.encode()
    repl['xl/styles.xml']=(O/'evidencia/estilos_v15_normalizados.xml').read_bytes()
    dst=O/NAME;assert not dst.exists()
    with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
        for n in z.namelist():out.writestr(n,repl.get(n,z.read(n)))
    zz=zipfile.ZipFile(dst);assert zz.testzip() is None
    for n in zz.namelist():
        if n.endswith(('.xml','.rels')):ET.fromstring(zz.read(n))
    wf=openpyxl.load_workbook(dst,read_only=True,data_only=False);wv=openpyxl.load_workbook(dst,read_only=True,data_only=True)
    checks=[]
    for e in m['edits']:
        c=wf[e['sheet']][e['cell']];v=wv[e['sheet']][e['cell']].value;want=e['expected']
        assert c.value==e.get('formula',e.get('value')),(e['cell'],c.value)
        assert abs(v-want)<1e-12 if isinstance(want,(float,int)) else v==want,(e['cell'],v,want)
        checks.append(dict(hoja=e['sheet'],celda=e['cell'],esperado=want,observado=v))
    prior=openpyxl.load_workbook(SRC,read_only=True,data_only=False);allow={(e['sheet'],e['cell']) for e in m['edits']};diffs=[]
    for sn in names:
        for rb,ra in zip(prior[sn],wf[sn]):
            for b,c in zip(rb,ra):
                if b.value!=c.value:
                    assert (sn,c.coordinate) in allow,(sn,c.coordinate)
                    diffs.append(dict(hoja=sn,celda=c.coordinate,antes=b.value,despues=c.value))
    for s,c,v in [('18_Foto_Intertrimestral','G8',1023),('18_Foto_Intertrimestral','D8',1150),('18_Foto_Intertrimestral','E8',611),('26_Avance_Curricular_Resumen','C1853',1150)]:assert wv[s][c].value==v
    charts=[n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n)];tables=[n for n in z.namelist() if re.fullmatch(r'xl/tables/table\d+\.xml',n)]
    assert len(charts)==13 and len(tables)==35
    assert all(z.read(n)==zz.read(n) for n in charts+tables)
    changed=[n for n in z.namelist() if z.read(n)!=zz.read(n)];assert set(changed)==set(repl)
    for p,h in m['fuentes'].items():assert sha(p)==h['sha256']
    assert git()==m['git_antes']
    dump('QA_ESTRUCTURAL_CALCULOS.json',dict(estado='OK_PENDIENTE_APERTURA_NATIVA',controles=checks,partes_modificadas=changed,hojas_conservadas_bytes=len(names)-len(by),graficos_conservados_bytes=13,tablas_conservadas_bytes=35,denominadores_conservados=True,fuentes_conservadas=True,git_antes=m['git_antes'],git_despues=git(),sha256=sha(dst)))
    dump('DIFF_CELDAS_EXISTENTES.json',diffs)
    print(str(dst))
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('fase',choices=['prepare','finish']);args=p.parse_args();(prepare if args.fase=='prepare' else finish)()
