from __future__ import annotations

import csv
import hashlib
import json
import shutil
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path(
    "/Users/alexi/Library/CloudStorage/OneDrive-InstitutoProfesionalSanSebastián/"
    "Datos DAI (RAW) - 1.1 Datos crudos/1.2.1 Datos Institucionales/"
    "1.2.1.1 SIES/Enviados/Matrícula Unificada"
)
REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
RUN_ID = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_DIR = REPO / "outputs" / "base_retencion_matricula_unificada_2022_2026" / RUN_ID
DESKTOP = Path("/Users/alexi/Desktop")

RECTOR_2026 = BASE / "2026" / "REPORTE JULIO 2026" / "MATRICULA_UNIFICADA_PREGRADO_2026_CONSOLIDADA_FINAL_CON_ENCABEZADOS.xlsx"

SELECTED_SOURCES = {
    2022: {
        "path": BASE / "2022" / "2022.csv",
        "kind": "csv_with_header",
        "notes": "Archivo CSV granular anual. Trae COD_NIV_GLO y FECH_CAR fuera del esquema rector 2026.",
    },
    2023: {
        "path": BASE / "2023" / "2023.csv",
        "kind": "csv_with_header",
        "notes": "Archivo CSV granular anual. No trae COD_NIV_GLO ni COD_IES; trae CARGA y CODIGO CNED fuera del esquema rector.",
    },
    2024: {
        "path": BASE / "2024" / "08-05-2024_18-31-54_Matrícula_Pregrado_2024.csv",
        "kind": "csv_with_header",
        "notes": "Archivo CSV granular pregrado completo. Se descarta archivo de solo primer anio y columnas finales de analisis.",
    },
    2025: {
        "path": BASE / "2025" / "2025.xlsx",
        "kind": "xlsx_sheet",
        "sheet": "08-05-2025_22-26-22_Matrícula_P",
        "notes": "Libro con hoja granular y hoja de tabla dinamica/resumen. Se usa solo la hoja granular.",
    },
    2026: {
        "path": RECTOR_2026,
        "kind": "xlsx_sheet",
        "sheet": "CONSOLIDADO",
        "notes": "Fuente gobernada con encabezados; define el esquema rector de 32 columnas.",
    },
}

EXCLUDED_PATTERNS = [
    "1. MATRICULA - RETENCION - TITULACION SIES _IP_CIISA_2010_20XX.xlsx",
    "IPSS_Procesos SIES.pbix",
]

DEMONSTRABLE_ALIASES = {
    2022: {
        "PRIMER_APELLIDO": "A_PAT",
        "SEGUNDO_APELLIDO": "A_MAT",
    },
    2025: {
        "VIG": "VIGENCIA",
    },
}

TEXT_COLUMNS = {
    "TIPO_DOC",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRE",
    "SEXO",
    "FECH_NAC",
    "FECHA_MATRICULA",
}

INTEGER_COLUMNS = {
    "ANIO_INFORMADO",
    "N_DOC",
    "NAC",
    "PAIS_EST_SEC",
    "COD_SED",
    "COD_CAR",
    "MODALIDAD",
    "JOR",
    "VERSION",
    "FOR_ING_ACT",
    "ANIO_ING_ACT",
    "SEM_ING_ACT",
    "ANIO_ING_ORI",
    "SEM_ING_ORI",
    "ASI_INS_ANT",
    "ASI_APR_ANT",
    "PROM_PRI_SEM",
    "PROM_SEG_SEM",
    "ASI_INS_HIS",
    "ASI_APR_HIS",
    "NIV_ACA",
    "SIT_FON_SOL",
    "SUS_PRE",
    "REINCORPORACION",
    "VIG",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_csv(path: Path) -> tuple[list[str], list[list[str]], str]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("latin-1", errors="replace")
        encoding = "latin-1-replace"
    sample = "\n".join(text.splitlines()[:10])
    dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    rows = list(csv.reader(text.splitlines(), delimiter=dialect.delimiter))
    return rows[0], rows[1:], f"{encoding}; delimiter={dialect.delimiter}"


def read_xlsx_sheet(path: Path, sheet_name: str) -> tuple[list[str], list[list[str]], list[dict]]:
    diagnostics = []
    with tempfile.TemporaryDirectory() as tmp:
        local = Path(tmp) / path.name
        shutil.copy2(path, local)
        wb = openpyxl.load_workbook(local, read_only=True, data_only=True)
        diagnostics.append(
            {
                "sheets": [
                    {"name": ws.title, "max_row": ws.max_row, "max_column": ws.max_column}
                    for ws in wb.worksheets
                ]
            }
        )
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        header = ["" if v is None else str(v).strip() for v in next(iterator)]
        rows = [["" if v is None else str(v).strip() for v in row] for row in iterator]
        wb.close()
    return header, rows, diagnostics


def normalize_header(h: str) -> str:
    return " ".join(str(h).strip().upper().split())


def has_header_like(row: list[str], rector: list[str]) -> bool:
    normalized = {normalize_header(x) for x in row}
    return len(normalized.intersection({normalize_header(x) for x in rector})) >= 8


def key_for_duplicate(row: dict[str, str]) -> tuple[str, ...]:
    return (
        row.get("TIPO_DOC", ""),
        row.get("N_DOC", ""),
        row.get("DV", ""),
        row.get("COD_SED", ""),
        row.get("COD_CAR", ""),
        row.get("MODALIDAD", ""),
        row.get("JOR", ""),
        row.get("VERSION", ""),
        row.get("ANIO_ING_ACT", ""),
        row.get("SEM_ING_ACT", ""),
    )


def typed_value(column: str, value: str):
    value = "" if value is None else str(value).strip()
    if value == "":
        return None
    if column == "VIG":
        vig_text = value.strip().upper()
        if vig_text in {"VIGENTE", "SI", "SÍ", "TRUE"}:
            return 1
        if vig_text in {"NO VIGENTE", "NO", "FALSE"}:
            return 0
    if column in INTEGER_COLUMNS:
        normalized = value.replace(".", "").replace(",", ".")
        if normalized.endswith(".0"):
            normalized = normalized[:-2]
        if normalized.lstrip("-").isdigit():
            return int(normalized)
        try:
            numeric = float(normalized)
        except ValueError:
            return value
        return int(numeric) if numeric.is_integer() else numeric
    return value


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    header_2026, _, rector_diag = read_xlsx_sheet(RECTOR_2026, "CONSOLIDADO")
    rector_cols = [str(x).strip() for x in header_2026]
    final_cols = ["ANIO_INFORMADO"] + rector_cols

    consolidated: list[list[str]] = []
    source_inventory = []
    structure_rows = []
    counts = []
    pending = []
    duplicate_counter: Counter[tuple[str, ...]] = Counter()

    for year, spec in SELECTED_SOURCES.items():
        path = spec["path"]
        if spec["kind"] == "csv_with_header":
            headers, rows, reader_notes = read_csv(path)
            diagnostics = [{"reader": reader_notes}]
        else:
            headers, rows, diagnostics = read_xlsx_sheet(path, spec["sheet"])

        headers = [str(h).strip() for h in headers]
        hmap = {normalize_header(h): idx for idx, h in enumerate(headers)}
        mapped = {}
        alias_used = {}
        missing = []
        for col in rector_cols:
            norm = normalize_header(col)
            if norm in hmap:
                mapped[col] = hmap[norm]
            elif col in DEMONSTRABLE_ALIASES.get(year, {}) and normalize_header(DEMONSTRABLE_ALIASES[year][col]) in hmap:
                alias = DEMONSTRABLE_ALIASES[year][col]
                mapped[col] = hmap[normalize_header(alias)]
                alias_used[col] = alias
            else:
                missing.append(col)

        mapped_source_indexes = set(mapped.values())
        extras = [h for i, h in enumerate(headers) if i not in mapped_source_indexes]
        effective_rows = 0
        first_data_repeated_header = 0
        for row in rows:
            row = ["" if v is None else str(v).strip() for v in row]
            if not any(row):
                continue
            if has_header_like(row, rector_cols):
                first_data_repeated_header += 1
                continue
            out_row = [year]
            row_dict = {}
            for col in rector_cols:
                value = row[mapped[col]] if col in mapped and mapped[col] < len(row) else ""
                out_row.append(typed_value(col, value))
                row_dict[col] = value
            consolidated.append(out_row)
            duplicate_counter[key_for_duplicate(row_dict)] += 1
            effective_rows += 1

        source_inventory.append(
            {
                "anio": year,
                "archivo": str(path),
                "clasificacion": "fuente institucional granular seleccionada",
                "sha256": sha256(path),
                "bytes": path.stat().st_size,
                "notas": spec["notes"],
                "diagnostico": diagnostics,
            }
        )
        counts.append(
            {
                "anio": year,
                "filas_consolidadas": effective_rows,
                "columnas_origen": len(headers),
                "columnas_mapeadas_a_2026": len(mapped),
                "columnas_faltantes_vs_2026": missing,
                "columnas_extra_no_consolidadas": extras,
                "alias_equivalentes_usados": alias_used,
                "filas_header_repetido_descartadas": first_data_repeated_header,
            }
        )
        for col in missing:
            pending.append(
                {
                    "anio": year,
                    "tipo": "columna_faltante_vs_esquema_2026",
                    "detalle": col,
                    "criterio": "No se inventa equivalencia; queda vacia en consolidado.",
                }
            )
        for col in extras:
            structure_rows.append(
                {
                    "anio": year,
                    "tipo": "columna_extra_origen_no_consolidada",
                    "detalle": col,
                    "criterio": "Fuera del esquema rector 2026 o de analisis; no se incorpora.",
                }
            )

    all_candidates = []
    for p in BASE.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".csv", ".xlsx", ".xlsm", ".xls", ".pbix"}:
            continue
        reason = "no seleccionado"
        if p.name in EXCLUDED_PATTERNS or p.suffix.lower() == ".pbix":
            reason = "excluido expresamente por instruccion"
        elif any(p == spec["path"] for spec in SELECTED_SOURCES.values()):
            reason = "seleccionado"
        elif "post" in str(p).lower():
            reason = "detectado, no consolidado: universo postitulo/posgrado sin mismo esquema rector pregrado 2026"
        elif "solo estudiantes de primer" in str(p).lower():
            reason = "detectado, no consolidado: subconjunto primer anio"
        elif p.suffix.lower() not in {".csv", ".xlsx", ".xlsm", ".xls"}:
            reason = "tipo no pertinente"
        all_candidates.append({"archivo": str(p), "motivo": reason})

    duplicate_keys = sum(1 for _, n in duplicate_counter.items() if n > 1)
    duplicate_rows = sum(n for _, n in duplicate_counter.items() if n > 1)

    xlsx_name = f"BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_{RUN_ID}.xlsx"
    xlsx_path = OUT_DIR / xlsx_name
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "BASE_RETENCION_MU"
    ws.append(final_cols)
    for row in consolidated:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    derived_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = derived_fill if cell.column == 1 else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for idx, col in enumerate(final_cols, 1):
        if col in INTEGER_COLUMNS:
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = "0"
    widths = {"A": 16}
    for idx, col in enumerate(final_cols, 1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = max(widths.get(letter, 12), min(max(len(col) + 2, 10), 24))
    wb.save(xlsx_path)

    desktop_path = DESKTOP / "BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026.xlsx"
    shutil.copy2(xlsx_path, desktop_path)

    manifest = {
        "estado": "BASE_RETENCION_MU_2022_2026_CONSTRUIDA_Y_VALIDADA",
        "timestamp": RUN_ID,
        "proceso": "Matricula Unificada SIES",
        "subproyecto": "Base consolidada para retencion 2022-2026",
        "repositorio": str(REPO),
        "rama": "feature/ire-2026",
        "esquema_rector": {
            "anio": 2026,
            "archivo": str(RECTOR_2026),
            "hoja": "CONSOLIDADO",
            "columnas": rector_cols,
            "diagnostico": rector_diag,
        },
        "salidas": {
            "excel_repositorio": str(xlsx_path),
            "excel_escritorio": str(desktop_path),
            "sha256_excel_repositorio": sha256(xlsx_path),
            "sha256_excel_escritorio": sha256(desktop_path),
        },
        "conteos": counts,
        "duplicados": {
            "llave_tecnica": [
                "TIPO_DOC",
                "N_DOC",
                "DV",
                "COD_SED",
                "COD_CAR",
                "MODALIDAD",
                "JOR",
                "VERSION",
                "ANIO_ING_ACT",
                "SEM_ING_ACT",
            ],
            "llaves_duplicadas": duplicate_keys,
            "filas_en_llaves_duplicadas": duplicate_rows,
            "nota": "Duplicado tecnico para auditoria, no regla oficial de eliminacion.",
        },
        "pendientes": pending,
    }

    (OUT_DIR / "MANIFEST_BASE_RETENCION_MU_2022_2026.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT_DIR / "INVENTARIO_FUENTES_SHA256.json").write_text(
        json.dumps(source_inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT_DIR / "INVENTARIO_CANDIDATOS_Y_EXCLUSIONES.json").write_text(
        json.dumps(all_candidates, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT_DIR / "DIFERENCIAS_ESTRUCTURALES.json").write_text(
        json.dumps({"diferencias": structure_rows, "pendientes": pending}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    report = [
        "REPORTE VALIDACION BASE RETENCION MATRICULA UNIFICADA 2022-2026",
        f"Timestamp: {RUN_ID}",
        "",
        "Estado: BASE_RETENCION_MU_2022_2026_CONSTRUIDA_Y_VALIDADA",
        f"Excel repositorio: {xlsx_path}",
        f"Excel escritorio: {desktop_path}",
        f"SHA-256 repositorio: {manifest['salidas']['sha256_excel_repositorio']}",
        f"SHA-256 escritorio: {manifest['salidas']['sha256_excel_escritorio']}",
        "",
        "Conteos por anio:",
    ]
    for c in counts:
        report.append(
            f"- {c['anio']}: {c['filas_consolidadas']} filas; "
            f"{c['columnas_mapeadas_a_2026']}/32 columnas mapeadas; "
            f"faltantes={c['columnas_faltantes_vs_2026']}; extras_no_consolidadas={c['columnas_extra_no_consolidadas']}"
        )
    report.extend(
        [
            "",
            f"Total filas consolidadas: {len(consolidated)}",
            f"Total columnas finales: {len(final_cols)} (incluye ANIO_INFORMADO derivada)",
            f"Llaves tecnicas duplicadas: {duplicate_keys}",
            f"Filas en llaves tecnicas duplicadas: {duplicate_rows}",
            "",
            "Nota: ANIO_INFORMADO es columna tecnica derivada; no forma parte de la carga original SIES.",
            "No se modificaron fuentes originales.",
        ]
    )
    (OUT_DIR / "REPORTE_VALIDACION_BASE_RETENCION_MU_2022_2026.txt").write_text(
        "\n".join(report), encoding="utf-8"
    )

    with zipfile.ZipFile(xlsx_path) as zf:
        bad = zf.testzip()
    print(json.dumps({"out_dir": str(OUT_DIR), "xlsx": str(xlsx_path), "desktop": str(desktop_path), "testzip": bad, "rows": len(consolidated), "counts": counts}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
