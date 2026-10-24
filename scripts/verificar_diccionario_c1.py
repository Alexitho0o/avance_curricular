#!/usr/bin/env python3
"""Verifica datos, fórmulas y sectores de la edición C1 sin escribir XLSX."""
import json,hashlib
from pathlib import Path
from collections import Counter
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'data/diccionario_c1/20261016'
OUT=ROOT/'outputs/diccionario_c1_20261016/DICCIONARIO_CRUCE_SIES_UMAS_2007_2026_C1.xlsx'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
    payload=json.loads((BASE/'libro_c1.json').read_text())
    source=openpyxl.load_workbook(payload['source'],data_only=False)
    target=openpyxl.load_workbook(OUT,data_only=False)
    cached=openpyxl.load_workbook(OUT,data_only=True)
    assert sha(payload['source_original'])==payload['source_sha256']
    assert sha(payload['source'])==payload['source_sha256']
    differences=[];formats=Counter();checked=0;scroll=[]
    for s in source:
        t=target[s.title]
        for row in s:
            for cell in row:
                if cell.value is None:continue
                checked+=1;other=t[cell.coordinate]
                if cell.value!=other.value:differences.append([s.title,cell.coordinate,str(cell.value),str(other.value)])
                for attr in ('font','fill','alignment','number_format','border','protection'):
                    if str(getattr(cell,attr))!=str(getattr(other,attr)):formats[attr]+=1
        # topLeftCell is the scrolled viewport, not the frozen row/column count.
        ap=s.sheet_view.pane;bp=t.sheet_view.pane
        assert (ap is None)==(bp is None),(s.title,'paneles')
        if ap is not None:
            assert (ap.xSplit,ap.ySplit,ap.state)==(bp.xSplit,bp.ySplit,bp.state),(s.title,'paneles')
            if ap.topLeftCell!=bp.topLeftCell:scroll.append({'hoja':s.title,'antes':ap.topLeftCell,'ahora':bp.topLeftCell})
        assert set(s.tables)==set(t.tables),(s.title,'tablas')
        for name,table in s.tables.items():
            # Worksheet tables may expose references instead of objects through items().
            a=s.tables[name];b=t.tables[name]
            assert a.ref==b.ref and a.autoFilter==b.autoFilter,(s.title,name,'filtro')
        assert {i for i,v in s.row_dimensions.items() if v.hidden}=={i for i,v in t.row_dimensions.items() if v.hidden},(s.title,'filas ocultas')
    assert not differences,differences[:5]
    assert not formats,dict(formats)
    protected=json.loads((ROOT/'data/unificacion_proyectos/20261015/control_preservacion_final.json').read_text())['referencias_sha256']
    for name,digest in protected.items():assert sha(ROOT/name)==digest,('referencia protegida',name)
    s=source['Diccionario SIES'];t=target['Diccionario SIES']
    headers=[c.value for c in s[5]]
    columns=[headers.index(x)+1 for x in ('Código Único','Sector IPSS 2026','Estado sector')]
    a=[[s.cell(i,j).value for j in columns] for i in range(6,147)]
    b=[[t.cell(i,j).value for j in columns] for i in range(6,147)]
    assert a==b
    newchecked=0
    for info in payload['sheets']:
        sheet=target[info['name']]
        assert [sheet.cell(5,j+1).value for j in range(len(info['headers']))]==info['headers']
        for i,row in enumerate(info['rows'],6):
            for j,value in enumerate(row,1):
                actual=sheet.cell(i,j).value
                assert (None if actual=='' else actual)==(None if value=='' else value),(info['name'],i,j,actual,value)
                newchecked+=1
    errors=[]
    for sheet in cached:
        for row in sheet:
            for c in row:
                if c.data_type=='e':errors.append((sheet.title,c.coordinate,c.value))
    assert not errors,errors[:10]
    assert cached['Consulta B5']['B20'].value=='CORROBORADA_EN_RESERVA'
    assert [cached['Consulta B5'][f'B{i}'].value for i in range(35,39)]==[197,211,15,1]
    summary={'archivo':str(OUT),'sha256':sha(OUT),'hojas':target.sheetnames,'celdas_originales_verificadas':checked,
             'celdas_nuevas_verificadas':newchecked,'sectores_por_codigo_preservados':len(a),
             'diferencias_formato_serializado':dict(formats),'posicion_desplazamiento_reiniciada':scroll,'errores_formula':errors,
             'fuente_usuario_intacta':True,'tablas_filtros_paneles_y_filas_ocultas_preservados':True,
             'referencias_protegidas_intactas':len(protected),
             'recalculo_nativo_excel_verificado':False}
    (BASE/'qa_xlsx.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    files=[BASE/'libro_c1.json',BASE/'fuente_usuario.xlsx',BASE/'autor/generar_c1.mjs',Path(__file__),BASE/'qa_formulas.json',BASE/'qa_xlsx.json',ROOT/'data/matriz_academica_b4/20261016/manifest.json',ROOT/'data/reglas_historicas_b5/20261016/manifest.json',OUT]
    (BASE/'entrega.json').write_text(json.dumps({'fase':'C1','estado':'EDICION_SEPARADA_VERIFICADA','archivo':str(OUT),'sha256':sha(OUT),'original_preservado':payload['source_original'],'fuente_sha256':payload['source_sha256'],'archivos_sha256':{str(p.relative_to(ROOT)):sha(p) for p in files}},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(summary,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
