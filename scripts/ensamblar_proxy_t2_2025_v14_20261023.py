"""C: prepara una edición acotada o ensambla sus celdas recalculadas por Artifact.

Nunca escribe fuentes. Solo exporta agregados y documentación; conserva formatos,
tablas y gráficos originales. Identificadores se usan solo en memoria para QA.
"""
import argparse
from pathlib import Path
import json, hashlib, re, zipfile, shutil, subprocess
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape
import openpyxl

R = Path(__file__).resolve().parents[1]
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
Q = lambda t: '{'+NS+'}'+t
SHA = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()

def git():
    return {k:subprocess.check_output(['git',*args],cwd=R,text=True).strip()
            for k,args in [('rama',['branch','--show-current']),('HEAD',['rev-parse','HEAD'])]}

def dump(p,x):
    p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')

def prepare(o):
    assert not (o/'EDICION.json').exists(), 'No sobrescribir edición'
    result=json.loads((o/'RESULTADO_AGREGADO.json').read_text())
    t=result['totales']
    assert (t['denom'],t['vigR'],t['vigP'],t['observados'],t['numerador_propuesto'])==(1150,600,11,694,611)
    for p,h in result['fuentes'].items(): assert SHA(p)==h['sha256']
    src=R/'outputs/seguimiento_t1_hasta_2026_20261023_v13/RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V13.xlsx'
    assert SHA(src)=='c4f506adcf16e9ae2bbf072d21be5033e30ee6ee81eb5bc87171a73867cc164b'
    w=openpyxl.load_workbook(src,read_only=True,data_only=False)
    edits=[]
    def add(s,c,value=None,formula=None,expected=None):
        assert w[s][c].value is not None, (s,c)
        e={'sheet':s,'cell':c,'before':w[s][c].value}
        if formula is None:e.update(value=value,expected=value)
        else:e.update(formula=formula,expected=expected)
        edits.append(e)
    s26='26_Avance_Curricular_Resumen';s18='18_Foto_Intertrimestral';s52='52_Codigos_SIES_Revision';s53='53_Codigos_T1_Confirmados';s28='28_QA_Matematica'
    add(s26,'D1852',694)
    add(s26,'E1852',formula=f'=SUMIFS(\'{s53}\'!$G$6:$G$223,\'{s53}\'!$A$6:$A$223,2025,\'{s53}\'!$C$6:$C$223,2)',expected=611)
    add(s26,'F1852',formula=w[s26]['F1852'].value,expected=611/1150)
    add(s26,'H1852','C · Incluye R/P; proxy, no foto de cierre.')
    add(s18,'E8',formula=f"='{s26}'!E1852",expected=611)
    add(s18,'F8',formula='=IF(D8=0,"",E8/D8)',expected=611/1150)
    add(s28,'C11',611)
    add(s28,'D11',formula=f"='{s26}'!E1852",expected=611)
    add(s28,'E11',formula='=D11',expected=611)
    add(s28,'F11',formula='=IF(C11=E11,"OK","FALLIDO")',expected='OK')
    add(s28,'H11','C · RUT+DV y puente P: 600 R + 11 P.')
    add(s28,'C12',611/1150)
    add(s28,'D12',formula=f"='{s26}'!F1852",expected=611/1150)
    add(s28,'E12',formula='=D12',expected=611/1150)
    add(s28,'F12',formula='=IF(ROUND(C12,10)=ROUND(E12,10),"OK","FALLIDO")',expected='OK')
    add(s52,'A1','52 · Seguimiento de T1 en T2 y T3 hasta 2026\nC/D · Se conserva el código T1, las bases y los denominadores. V14 corrige el proxy T2-2025 por identidad completa, incluidos documentos P. El seguimiento académico se separa de la revisión del código SIES.')
    add(s52,'E8',formula=w[s52]['E8'].value,expected=611)
    add(s52,'G8',formula=w[s52]['G8'].value,expected=611/1150)
    add(s52,'I8','B/C · 611 por cruce R/P; desglose conciliado con hoja 53')
    add(s52,'A14','C · T2-2025: 611/1.150; incluye 11 documentos P verificados. E · No se recuperó el cálculo original de 600. El estado en PROMEDIOS no acredita vigencia de matrícula al cierre.')
    add(s52,'A32','C · Se conserva el diccionario v7 y la revisión de los períodos 2 y 3. V14 corrige únicamente el proxy T2-2025 y sus referencias en hojas 18, 26, 28, 52 y 53. Las otras 49 hojas, los denominadores y los 13 gráficos se conservan.')
    groups={d['codigo']:d for d in result['por_codigo_T1']}
    count=0
    for row in range(6,224):
        if (w[s53].cell(row,1).value,w[s53].cell(row,3).value)!=(2025,2): continue
        d=groups[w[s53].cell(row,5).value]
        assert d['nombre']==w[s53].cell(row,4).value
        assert d['denom']==w[s53].cell(row,6).value==w[s53].cell(row,8).value
        add(s53,f'G{row}',d['numerador_propuesto'])
        add(s53,f'I{row}',formula=w[s53][f'I{row}'].value,expected=d['tasa_propuesta'])
        add(s53,f'J{row}','C · Actividad 2025-T2 y estado VIGENTE; R/P. No foto de cierre.')
        count+=1
    assert count==36
    w.close()
    dump(o/'EDICION.json',{'source':str(src),'source_sha256':SHA(src),'scratch':str(Path('/tmp/retencion_proxy_t2_v14_20261023')),'edits':edits,'git_antes':git(),'fuentes':result['fuentes']})
    print(json.dumps({'ediciones':len(edits),'grupos':count,'denominador':1150,'numerador':611}))

def finish(o):
    m=json.loads((o/'EDICION.json').read_text());src=Path(m['source'])
    assert SHA(src)==m['source_sha256'];assert git()==m['git_antes']
    tmp=Path(m['scratch'])/'artifact_recalculado.xlsx'
    wf=openpyxl.load_workbook(tmp,read_only=True,data_only=False)
    wv=openpyxl.load_workbook(tmp,read_only=True,data_only=True)
    for e in m['edits']:
        f=wf[e['sheet']][e['cell']];v=wv[e['sheet']][e['cell']].value
        assert f.value==e.get('formula',e.get('value')),(e['sheet'],e['cell'],'contenido')
        want=e['expected']
        assert (abs(v-want)<1e-12 if isinstance(v,(int,float)) and isinstance(want,(int,float)) else v==want),(e['sheet'],e['cell'],v,want)
    wf.close();wv.close()
    z=zipfile.ZipFile(src);a=zipfile.ZipFile(tmp)
    names=[s.get('name') for s in ET.fromstring(z.read('xl/workbook.xml')).find(Q('sheets'))]
    anames=[s.get('name') for s in ET.fromstring(a.read('xl/workbook.xml')).find(Q('sheets'))]
    assert names==anames
    strings=[''.join(n.itertext()) for n in ET.fromstring(a.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in a.namelist() else []
    replacement={};byname={}
    for e in m['edits']:byname.setdefault(e['sheet'],[]).append(e)
    for name,edits in byname.items():
        part=f'xl/worksheets/sheet{names.index(name)+1}.xml'
        txt=z.read(part).decode();tree=ET.fromstring(a.read(part));cells={c.get('r'):c for c in tree.iter(Q('c'))}
        for e in edits:
            address=e['cell'];c=cells[address]
            pat=rf'<c\b(?=[^>]*\br="{address}")[^>]*>.*?</c>'
            hit=re.search(pat,txt,re.S);assert hit,(name,address)
            old=ET.fromstring(hit.group())
            attrs=dict(old.attrib);attrs.pop('t',None)
            if c.get('t'):attrs['t']=c.get('t')
            if attrs.get('t')=='s':
                st=strings[int(c.find(Q('v')).text)];attrs['t']='inlineStr';body='<is><t xml:space="preserve">'+escape(st)+'</t></is>'
            else:
                body=''.join(ET.tostring(n,encoding='unicode').replace('ns0:','').replace(':ns0','') for n in c)
                body=re.sub(r' xmlns="[^"]+"','',body)
            new='<c '+ ' '.join(k+'="'+escape(v,{'"':'&quot;'})+'"' for k,v in attrs.items())+'>'+body+'</c>'
            txt=txt[:hit.start()]+new+txt[hit.end():]
        replacement[part]=txt.encode()
    assert len(replacement)==5
    assert 'xl/calcChain.xml' not in z.namelist(),'Cadena de cálculo requiere tratamiento explícito'
    dst=o/'RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V14.xlsx'
    assert not dst.exists()
    with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
        for n in z.namelist():out.writestr(n,replacement.get(n,z.read(n)))
    zz=zipfile.ZipFile(dst);assert zz.testzip() is None
    changed=[n for n in z.namelist() if z.read(n)!=zz.read(n)]
    assert set(changed)==set(replacement)
    for n in zz.namelist():
        if n.endswith(('.xml','.rels')):ET.fromstring(zz.read(n))
    final=openpyxl.load_workbook(dst,read_only=True,data_only=False)
    prior=openpyxl.load_workbook(src,read_only=True,data_only=False)
    cached=openpyxl.load_workbook(dst,read_only=True,data_only=True)
    allow={(e['sheet'],e['cell']) for e in m['edits']}; diffs=[]
    for name in names:
        for rows in zip(prior[name],final[name]):
            for b,c in zip(*rows):
                assert getattr(b,'_style_id',0)==getattr(c,'_style_id',0) and b.number_format==c.number_format
                if b.value!=c.value:
                    assert (name,c.coordinate) in allow,(name,c.coordinate)
                    diffs.append({'hoja':name,'celda':c.coordinate,'antes':b.value,'despues':c.value})
    for e in m['edits']:
        got=cached[e['sheet']][e['cell']].value;want=e['expected']
        assert (abs(got-want)<1e-12 if isinstance(got,(int,float)) and isinstance(want,(int,float)) else got==want)
    prior_cached=openpyxl.load_workbook(src,read_only=True,data_only=True)
    for row in range(6,12):
        if row!=8:
            for c in 'DEFG':assert cached['52_Codigos_SIES_Revision'][f'{c}{row}'].value==prior_cached['52_Codigos_SIES_Revision'][f'{c}{row}'].value
    prior_cached.close()
    charts=[n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n)]
    tables=[n for n in z.namelist() if re.fullmatch(r'xl/tables/table\d+\.xml',n)]
    assert len(charts)==13 and len(tables)==35
    assert all(z.read(n)==zz.read(n) for n in charts+tables)
    for p,h in m['fuentes'].items():assert SHA(p)==h['sha256']
    origins=json.loads(Path(next(p for p in m['fuentes'] if p.endswith('ORIGENES_T1.json'))).read_text())
    individual={str(d[k]) for d in origins for k in ['persona','documento_original'] if d.get(k)}
    assert not any(str(e.get('value','')) in individual for e in m['edits'])
    desktop=Path.home()/'Desktop'/dst.name;assert not desktop.exists();shutil.copy2(dst,desktop);assert SHA(desktop)==SHA(dst)
    qa={'estado':'OK_CON_LIMITES_E','numerador_antes':600,'numerador_despues':611,'denominador_antes_y_despues':1150,'tasa_antes':600/1150,'tasa_despues':611/1150,'diferencia_pp':100*11/1150,'presencias':694,'cambios_esperados_con_cache':len(m['edits']),'celdas_contenido_modificado':len(diffs),'hojas_modificadas':list(byname),'otras_hojas_conservadas_bytes':49,'graficos_conservados_bytes':13,'tablas_conservadas_bytes':35,'estilos_conservados':True,'fuentes_conservadas':True,'denominadores_conservados':True,'nuevas_asignaciones_SIES':0,'git_antes':m['git_antes'],'git_despues':git(),'commit':False,'push':False,'limites':['Recálculo Artifact Tool y contraste independiente; no se ha validado apertura en Excel nativo.','El mecanismo original que precalculó600 permanece no localizado.','El indicador es proxy de actividad y estado del extracto, no matrícula vigente a fecha histórica.'],'archivo':str(dst),'sha256':SHA(dst),'escritorio':str(desktop)}
    dump(o/'QA_CORRECCION.json',qa);dump(o/'DIFF_CELDAS.json',diffs)
    manifest={'clasificacion':'B fuentes; C corrección de cruce; D autorización del usuario; E mecanismo original y vigencia histórica pendientes','proceso':'Retención intertrimestral IPSS162','alcance':'Solo proxy2025-T2 de cohorteT1-2025','autorizacion':'Corregir cruce si se confirma error, sin cambiar denominadores.','fuentes':m['fuentes'],'version_anterior':{'ruta':str(src),'sha256':m['source_sha256']},'resultado':qa,'archivos':{str(p.relative_to(R)):SHA(p) for p in [dst,o/'RESULTADO_AGREGADO.json',o/'EDICION.json',o/'QA_CORRECCION.json',o/'DIFF_CELDAS.json',R/'scripts/recalcular_proxy_t2_2025_v14_20261023.py',Path(__file__),R/'scripts/editar_proxy_t2_2025_v14_20261023.mjs']}}
    dump(o/'MANIFEST_CORRECCION.json',manifest)
    print(json.dumps(qa,ensure_ascii=False,indent=2))

if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('phase',choices=['prepare','finish']);ap.add_argument('output_dir',type=Path);args=ap.parse_args()
    (prepare if args.phase=='prepare' else finish)(args.output_dir.resolve())
