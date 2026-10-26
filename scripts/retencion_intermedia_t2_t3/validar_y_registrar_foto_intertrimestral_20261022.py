#!/usr/bin/env python3
"""Valida fuentes y registra la trazabilidad de la foto intertrimestral.

No modifica fuentes ni construye una salida con datos individuales. La inyección
de hojas debe ejecutarse sobre una copia local mediante el autor de hojas aprobado;
este script deja un registro verificable de entradas, hashes y estructura.
"""
from __future__ import annotations
import hashlib, json, sys
from pathlib import Path
from openpyxl import load_workbook

MANIFEST = Path(__file__).parents[2] / "procesos/retencion_intermedia_t2_t3/MANIFEST_RETENCION_FOTO_INTERTRIMESTRAL_20261022.json"

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()

def main() -> int:
    manifest = json.loads(MANIFEST.read_text())
    results = []
    for src in manifest["fuentes"]:
        p = Path(src["ruta"])
        item = {"rol": src["rol"], "ruta": str(p), "existe": p.exists()}
        if p.exists():
            item["sha256"] = sha256(p)
            if "sha256" in src:
                item["hash_esperado"] = src["sha256"]
                item["hash_ok"] = item["sha256"] == src["sha256"]
            try:
                wb = load_workbook(p, read_only=True, data_only=False)
                item["hojas"] = wb.sheetnames
                item["dimensiones"] = {s: [wb[s].max_row, wb[s].max_column] for s in wb.sheetnames}
            except Exception as exc:
                item["lectura_excel"] = f"ERROR: {exc}"
        results.append(item)
    out = {"manifest": str(MANIFEST), "fuentes": results, "clasificacion": "C", "nota": "Validación de insumos; no cambia originales."}
    print(json.dumps(out, ensure_ascii=False, indent=2))
    return 0 if all(x["existe"] and x.get("hash_ok", True) for x in results) else 2

if __name__ == "__main__":
    raise SystemExit(main())
