"""Lectura dirigida de preparación MU2025 para E04; no adjudica ni modifica fuentes."""
from pathlib import Path
import json,sys,collections,subprocess,hashlib
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado
read=lambda p:json.loads(Path(p).read_text())
e=next(x for x in read('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023/ERRORES6_DETALLE.json') if x['id']=='E04');person=e['persona'];codcli=e['CODCLI']
P=Path('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023/regla_e04');P.mkdir(exist_ok=True,mode=0o700)
assert subprocess.run(['git','-C',str(P),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
cfg=read(R/'outputs/origen_v1_inf35_03_20261023_150536/CONFIG.json');sources={p:h for p,h in cfg['fuentes_sha256'].items() if Path(p).suffix=='.xlsx' and Path(p).name!='GESTIONSGE_1703.xlsx'}
fields={'CODCARPR','CODCARR','CODCLI','CODIGOCARRERA','CODCARRERA','ANOMATRICULA','PERIODOMATRICULA','ANOINGRESO','PERIODOINGRESO','FECHAMATRICULA','JORNADA','REGIMEN','PLAN_DE_ESTUDIO','COD_IES','COD_SED','COD_CAR','MODALIDAD','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FECHA_MATRICULA','REINCORPORACION','VIG','ESTADOACADEMICO','PLANESTUDIO','PLAN_ESTUDIOS','NIVEL','CONTINUIDAD','ARREGLO','OBS','CARGA','FECH_CAR'}
records=[];inventory=[]
for p,hsh in sources.items():
 p=Path(p);assert sha256(p)==hsh
 wv=load_workbook(p,read_only=True,data_only=True);wf=load_workbook(p,read_only=True,data_only=False)
 for ws in wv:
  header=None
  for rn,row in enumerate(ws.iter_rows(max_row=min(15,ws.max_row),values_only=True),1):
   h=[str(v).strip().upper() if v is not None else '' for v in row]
   if 'N_DOC' in h or 'RUT' in h:header=h;start=rn;break
  if header is None:inventory.append({'fuente':str(p),'hoja':ws.title,'estado':'SIN_ENCABEZADO_IDENTIDAD_EN_15_FILAS'});continue
  inds=[i for i,h in enumerate(header) if h in ['N_DOC','RUT']];dv=next((i for i,h in enumerate(header) if h=='DV'),None);tipo=next((i for i,h in enumerate(header) if h=='TIPO_DOC'),None);matches=[]
  for rn,row in enumerate(ws.iter_rows(min_row=start+1,values_only=True),start+1):
   found=False;method=[]
   for i in inds:
    if i>=len(row):continue
    val=row[i]
    if header[i]=='N_DOC' and dv is not None:
     if tipo is not None and str(row[tipo]).upper() not in ['R','1']:continue
     val=str(int(val) if isinstance(val,(int,float)) else val)+'-'+str(row[dv]);m='N_DOC+DV'+('+TIPO_DOC' if tipo is not None else ' (tipo documental ausente)')
    else:m='RUT completo normalizado'
    try:key=rut_normalizado(val)
    except (ValueError,TypeError):continue
    if key==person:found=True;method.append(m)
   if not found:continue
   vals={get_column_letter(i+1)+':'+h:row[i] for i,h in enumerate(header) if h in fields and i<len(row)}
   rec={'caso':'E04','fuente':str(p),'sha256':hsh,'hoja':ws.title,'fila':rn,'metodo':method,'valores':vals,'formulas':{}};records.append(rec);matches.append(rec)
  if matches:
   targets={z['fila']:z for z in matches}
   for rn,row in enumerate(wf[ws.title].iter_rows(min_row=min(targets),max_row=max(targets),values_only=True),min(targets)):
    if rn not in targets:continue
    for i,h in enumerate(header):
     if h in fields and i<len(row) and isinstance(row[i],str) and row[i].startswith('='):targets[rn]['formulas'][get_column_letter(i+1)+':'+h]=row[i]
  inventory.append({'fuente':str(p),'hoja':ws.title,'encabezado':start,'filas_foco':len(matches)})
 wv.close();wf.close()
 print(p.name,'leído',flush=True)
(P/'PREPARACION_PRIVADA.json').write_text(json.dumps({'persona':person,'CODCLI':codcli,'records':records},ensure_ascii=False,indent=2,default=str))
public=[]
for r in records:
 vals={k:v for k,v in r['valores'].items() if k.split(':',1)[1] not in ['CODCLI','OBS','CARGA','ARREGLO']}
 public.append({k:v for k,v in r.items() if k!='valores'}|{'valores':vals})
result={'clasificacion':'B datos; C rastreo; E regla generadora no establecida por coincidencia','registros':public,'inventario':inventory,'fuentes':sources,'fuentes_intactas':all(sha256(p)==h for p,h in sources.items())}
text=json.dumps(result,ensure_ascii=False,indent=2,default=str);assert person not in text and codcli not in text
Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_preparacion.json').write_text(text)
print(json.dumps({'registros':len(records),'libros':len(sources),'con_version':[{'libro':Path(r['fuente']).name,'hoja':r['hoja'],'fila':r['fila'],'version':{k:v for k,v in r['valores'].items() if k.endswith(':VERSION')},'formula_version':{k:v for k,v in r['formulas'].items() if k.endswith(':VERSION')}} for r in public if any(k.endswith(':VERSION') for k in r['valores'])]},ensure_ascii=False))
