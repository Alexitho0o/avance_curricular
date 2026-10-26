"""Contraste agregado de versiones C1/J1 en Precarga MU2025 contra RAW2024."""
from pathlib import Path
import csv,json,sys,collections
from openpyxl import load_workbook
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado
prep=json.load(open('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_preparacion.json'));proc=next(Path(p) for p in prep['fuentes'] if Path(p).name=='IPSS_Procesamiento.xlsx');inc=json.load(open('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_incremental.json'));raw=next(Path(p) for p in inc['fuentes_sha256'] if '08-05-2024' in p)
try: rawtext=raw.read_text(encoding='utf-8-sig');encoding='utf-8-sig'
except UnicodeDecodeError: rawtext=raw.read_text(encoding='cp1252');encoding='cp1252'
rows=list(csv.DictReader(rawtext.splitlines(),delimiter=';'))
def val(x):return str(x if x is not None else '').strip().removesuffix('.0')
def identity(d):
 if d.get('TIPO_DOC') and val(d['TIPO_DOC'])!='R':return None
 try:return rut_normalizado(val(d.get('N_DOC'))+'-'+val(d.get('DV')))
 except (ValueError,TypeError):return None
old=collections.defaultdict(set)
for d in rows:
 p=identity(d)
 if p and val(d.get('COD_SED'))=='2' and val(d.get('COD_CAR'))=='1' and val(d.get('JOR'))=='1':old[p].add(val(d.get('VERSION')))
w=load_workbook(proc,read_only=True,data_only=True);s=w['Precarga'];hh=[val(x) for x in next(s.iter_rows(min_row=6,max_row=6,values_only=True))];out=collections.Counter();groups=collections.Counter();n=0
for rn,row in enumerate(s.iter_rows(min_row=7,values_only=True),7):
 d=dict(zip(hh,row));p=identity(d)
 if p and val(d.get('COD_SED'))=='2' and val(d.get('COD_CAR'))=='1' and val(d.get('JOR'))=='1':
  n+=1;v=val(d.get('VERSION'));oldvs=tuple(sorted(old.get(p,set())));out[(oldvs,v)]+=1;groups[(v,val(d.get('ANIO_ING_ACT')),val(d.get('FECH_CAR')))]+=1
w.close();result={'alcance':'C1/J1/sede2 en Precarga2025, identidad exacta RUT+DV+tipoR cuando presente; versiones previas del mismo código base2024. Descripción, no prueba de generador.','fuentes':{str(proc):sha256(proc),str(raw):sha256(raw)},'encoding_raw':encoding,'filas_precarga_c1j1':n,'transiciones':[{'versiones_2024':list(a),'version_2025_precarga':b,'filas':n} for (a,b),n in out.items()],'distribucion_lote_ingreso':[{'version':a,'ingreso_actual':b,'FECH_CAR':c,'filas':n} for (a,b,c),n in groups.items()]}
Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_grupo_precarga.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2))
