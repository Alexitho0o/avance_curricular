"""Inspecciona incremental MU2025 y estructura de preparación para E04; solo lectura."""
from pathlib import Path
import json,sys,csv,zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
R=Path('/Users/alexi/Documents/GitHub/avance_curricular');sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado
read=lambda p:json.loads(Path(p).read_text());e=next(x for x in read('/Users/alexi/.local/share/ipss/analisis_conflictos_22_20261023/ERRORES6_DETALLE.json') if x['id']=='E04');person=e['persona'];codcli=e['CODCLI'];prep=read('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_preparacion.json')
inv=read(R/'outputs/retencion_intermedia_t2_t3/20261021_093412_git_onedrive_mu2025/INVENTARIO_ONEDRIVE.json');files=[Path(x['ruta']) for x in inv['archivos'] if '/2025/Matricula Incremental/' in x['ruta'] and Path(x['ruta']).suffix in ['.csv','.xlsx']]
fields={'COD_IES','COD_SED','COD_CAR','MODALIDAD','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FECHA_MATRICULA','FECH_CAR','NIV_ACA','VIG'}
rows=[];inventory=[];hashes={}
for p in files:
 hashes[str(p)]=sha256(p)
 if p.suffix=='.xlsx':
  w=load_workbook(p,read_only=True,data_only=False);s=w['1'];header=[str(v) for v in next(s.iter_rows(values_only=True))];vals=list(s.iter_rows(min_row=2,values_only=True));w.close();datatype='formulas/original';start=2
 else:
  try:txt=p.read_text(encoding='utf-8-sig');enc='utf-8-sig'
  except UnicodeDecodeError:txt=p.read_text(encoding='cp1252');enc='cp1252'
  vals=list(csv.reader(txt.splitlines(),delimiter=';'))
  if 'N_DOC' in vals[0]:header=vals.pop(0);start=2
  else:
   xlsx=next(x for x in files if x.suffix=='.xlsx');w=load_workbook(xlsx,read_only=True,data_only=True);header=[str(v) for v in next(w['1'].iter_rows(values_only=True))][:32];w.close();assert all(len(r)==32 for r in vals);start=1
  datatype=enc
 count=0
 for rn,row in enumerate(vals,start):
  try:
   body=row[header.index('N_DOC')];dv=row[header.index('DV')];key=rut_normalizado(str(int(body))+'-'+str(dv))
  except (ValueError,TypeError,IndexError):continue
  if key!=person:continue
  if 'TIPO_DOC' in header:assert str(row[header.index('TIPO_DOC')]).strip().upper()=='R'
  count+=1;rows.append({'fuente':str(p),'hoja':'1' if p.suffix=='.xlsx' else 'CSV','fila':rn,'valores':{h:row[i] for i,h in enumerate(header) if h in fields and i<len(row)}})
 inventory.append({'fuente':str(p),'sha256':hashes[str(p)],'tipo_lectura':datatype,'filas_coincidentes':count})
proc=next(Path(p) for p in prep['fuentes'] if Path(p).name=='IPSS_Procesamiento.xlsx');w=load_workbook(proc,read_only=False,data_only=False);maps=[]
for rn,row in enumerate(w['Mapa'].iter_rows(values_only=True),1):
 if any('INFORMATICA' in str(v).upper() for v in row):maps.append({'hoja':'Mapa','fila':rn,'celdas':{get_column_letter(i+1):v for i,v in enumerate(row) if v is not None}})
cell=w['Precarga']['Q1025'];cellinfo={'hoja':'Precarga','celda':'Q1025','valor':cell.value,'tipo':cell.data_type,'formula':cell.data_type=='f','comentario':cell.comment.text if cell.comment else None,'hipervinculo':cell.hyperlink.target if cell.hyperlink else None}
formula_hist=[]
for s in w:
 for row in s.iter_rows():
  for c in row:
   if c.data_type=='f' and ('Mapa' in str(c.value) or 'VERSION' in str(c.value).upper()):formula_hist.append({'hoja':s.title,'celda':c.coordinate,'formula':c.value})
w.close()
meta=[]
for p in [proc]+[p for p in files if p.suffix=='.xlsx']:
 with zipfile.ZipFile(p) as z:
  core=ET.fromstring(z.read('docProps/core.xml')) if 'docProps/core.xml' in z.namelist() else None;dates={x.tag.split('}')[-1]:x.text for x in core if x.tag.split('}')[-1] in ['created','modified']} if core is not None else {};rels=[]
  for f in z.namelist():
   if f.startswith('xl/externalLinks/_rels/'):
    for rr in ET.fromstring(z.read(f)):rels.append(rr.attrib.get('Target'))
  meta.append({'fuente':str(p),'fechas_metadatos_no_fecha_cambio':dates,'vinculos_externos':rels,'conexiones':[n for n in z.namelist() if 'connections' in n.lower() or 'querytables' in n.lower() or 'vbaProject' in n]})
result={'incremental':rows,'inventario_incremental':inventory,'celda_precarga':cellinfo,'mapas':maps,'formulas_mapa_o_version':formula_hist,'metadatos':meta,'fuentes_sha256':hashes}
text=json.dumps(result,ensure_ascii=False,indent=2,default=str);assert person not in text and codcli not in text
Path('/Users/alexi/Documents/GitHub/avance_curricular/outputs/origen_versiones_c3j4_e04_20261023_163300/e04/regla_e04_incremental.json').write_text(text)
print(json.dumps({'incremental':rows,'celda_precarga':cellinfo,'cantidad_formulas_mapa_version':len(formula_hist),'metadatos':meta},ensure_ascii=False,default=str))
