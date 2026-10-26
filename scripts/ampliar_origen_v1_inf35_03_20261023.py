"""Inspecciona matrícula incremental y celdas/mapas preservando las fuentes."""
from pathlib import Path
import json,sys,csv,zipfile,xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado
O=R/'outputs/origen_v1_inf35_03_20261023_150536';read=lambda p:json.loads(Path(p).read_text());cfg=read(O/'CONFIG.json');P=Path(cfg['privado']);e=read(P/'EVIDENCIA.json');person=e['objetivos'][0]['persona'];codcli=e['objetivos'][0]['CODCLI']
def dump(p,x):
 p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str)+'\n')
 if p.is_relative_to(P):p.chmod(0o600)
inv=read(R/'outputs/retencion_intermedia_t2_t3/20261021_093412_git_onedrive_mu2025/INVENTARIO_ONEDRIVE.json');files=[Path(x['ruta']) for x in inv['archivos'] if '/2025/Matricula Incremental/' in x['ruta'] and Path(x['ruta']).suffix in ['.csv','.xlsx']]
fields={'COD_IES','COD_SED','COD_CAR','MODALIDAD','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FECHA_MATRICULA','FECH_CAR','NIV_ACA','VIG'}
rows=[];inventory=[];hashes={}
for p in files:
 hashes[str(p)]=sha256(p)
 if p.suffix=='.xlsx':
  w=load_workbook(p,read_only=True,data_only=False);s=w['1'];header=[str(v) for v in next(s.iter_rows(values_only=True))];vals=list(s.iter_rows(min_row=2,values_only=True));w.close();datatype='formulas/original';start=2
 else:
  try:txt=p.read_text(encoding='utf-8-sig');enc='utf-8-sig'
  except UnicodeDecodeError:txt=p.read_text(encoding='cp1252');enc='cp1252'
  vals=list(csv.reader(txt.splitlines(),delimiter=';'))
  if 'N_DOC' in vals[0]:header=vals.pop(0);start=2
  else:
   # Orden comprobado contra la hoja 1 del XLSX incremental; validación 32 columnas.
   xlsx=next(x for x in files if x.suffix=='.xlsx');w=load_workbook(xlsx,read_only=True,data_only=True);header=[str(v) for v in next(w['1'].iter_rows(values_only=True))][:32];w.close();assert all(len(r)==32 for r in vals);start=1
  datatype=enc
 count=0
 for rn,row in enumerate(vals,start):
  try:
   body=row[header.index('N_DOC')];dv=row[header.index('DV')];key=rut_normalizado(str(int(body))+'-'+str(dv))
  except (ValueError,TypeError,IndexError):continue
  if key!=person:continue
  count+=1;record={'fuente':str(p),'hoja':'1' if p.suffix=='.xlsx' else 'CSV','fila':rn,'valores':{h:row[i] for i,h in enumerate(header) if h in fields and i<len(row)}};rows.append(record)
 inventory.append({'fuente':str(p),'sha256':hashes[str(p)],'tipo_lectura':datatype,'filas_coincidentes':count})
# Mapas y propiedades de fuentes de preparación; no convertir fecha core.xml en fecha de cambio.
proc=next(Path(p) for p in cfg['fuentes_sha256'] if Path(p).name=='IPSS_Procesamiento.xlsx');w=load_workbook(proc,read_only=True,data_only=False);maps=[]
for rn,row in enumerate(w['Mapa'].iter_rows(values_only=True),1):
 if any('INFORMATICA' in str(v).upper() for v in row):maps.append({'hoja':'Mapa','fila':rn,'celdas':{get_column_letter(i+1):v for i,v in enumerate(row) if v is not None}})
# Celdas exactas de precarga: sin formulas/enlaces propios/comentarios; vínculos externos generales separados.
w.close();wn=load_workbook(proc,read_only=False,data_only=False);cell=wn['Precarga']['Q1023'];cellinfo={'hoja':'Precarga','celda':'Q1023','valor':cell.value,'tipo':cell.data_type,'formula':cell.data_type=='f','comentario':cell.comment.text if cell.comment else None,'hipervinculo':cell.hyperlink.target if cell.hyperlink else None};wn.close()
meta=[]
for p in [proc]+files:
 if p.suffix!='.xlsx':continue
 with zipfile.ZipFile(p) as z:
  core=ET.fromstring(z.read('docProps/core.xml')) if 'docProps/core.xml' in z.namelist() else None
  dates={x.tag.split('}')[-1]:x.text for x in core if x.tag.split('}')[-1] in ['created','modified']} if core is not None else {}
  rels=[]
  for f in z.namelist():
   if f.startswith('xl/externalLinks/_rels/'):
    for r in ET.fromstring(z.read(f)):rels.append(r.attrib.get('Target'))
  meta.append({'fuente':str(p),'fechas_metadatos_no_fecha_cambio':dates,'vinculos_externos':rels,'conexiones':[n for n in z.namelist() if 'connections' in n.lower() or 'querytables' in n.lower() or 'vbaProject' in n]})
master=R/'data/versiones_historicas_sies_umas/fase_5_20261015/maestro.json';hashes[str(master)]=sha256(master);offer=[x for x in read(master)['oferta_anual'] if x['codigo_unico'] in ['I162S2C1J1V1','I162S2C1J1V2','I162S2C1J2V4'] and str(x['anio']) in ['2024','2025']]
context=Path('/Users/alexi/.local/share/ipss/cierre_14_jornadas_15_20261023_133759/JORNADAS_15_CONTEXTO_COMPLETO.json');hashes[str(context)]=sha256(context);history=[x for x in read(context) if x.get('caso')=='INF35-03' and x.get('CODCLI')==codcli]
summary=[]
from collections import defaultdict,Counter
groups=defaultdict(list)
for x in history:
 if x.get('PLAN_DE_ESTUDIO') and str(x.get('ANO')) in ['2024','2025']:groups[(x['fuente'],str(x['ANO']),str(x.get('PERIODO')))].append(x)
for (source,year,period),a in sorted(groups.items()):summary.append({'fuente':source,'anio':year,'periodo':period,'filas':len(a),'planes':dict(Counter(x.get('PLAN_DE_ESTUDIO') for x in a)),'jornadas_exportacion_no_historico':dict(Counter(x.get('JORNADA') for x in a))})
dump(P/'AMPLIACION_PRIVADA.json',{'persona':person,'CODCLI':codcli,'historia_plan':history,'incremental':rows})
dump(O/'AMPLIACION.json',{'inventario_incremental':inventory,'filas_incremental':rows,'celda_precarga':cellinfo,'mapas':maps,'metadatos':meta,'oferta':offer,'historia_plan_resumen':summary,'fuentes_hashes_adicionales':hashes})
for p,h in hashes.items():assert sha256(p)==h
print(json.dumps({'incremental':rows,'celda_precarga':cellinfo,'mapas':maps,'metadatos':meta,'historia_plan_resumen':summary},ensure_ascii=False,default=str))
