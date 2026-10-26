"""C: reproduce la corrección del proxy T3-2025 sin modificar fuentes ni Excel.

Alcance cerrado: las 1150 personas de T1-2025, incluidas20 identidadesP con
puente explícito verificado. Este tratamiento no se extrapola a otrosP.
La evidencia del constructor anterior solo conserva valores precomputados;
no acredita cómo se calculó535 originalmente. Salida agregada, sin personas.
"""
from pathlib import Path
import argparse
import collections
import hashlib
import json
import re
import openpyxl


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,required=True,
                        help='Directorio para RESULTADO_AGREGADO.json; no sobrescribe una salida existente.')
    args=parser.parse_args()
    if (args.output_dir/'RESULTADO_AGREGADO.json').exists():
        parser.error('RESULTADO_AGREGADO.json ya existe; seleccione otro directorio.')
    R=Path('/Users/alexi/Documents/GitHub/avance_curricular')
    constructor_original=R/'outputs/correccion_proxy_t2_2025_20261023_v14/evidencia/add_promedios_sheet_original.py'
    sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
    assert sha(constructor_original)=='c09e9b3255cfb90f38ca1393cb49d62e8e061cf8c39fa081d8011525a70183d8'
    originpath=Path('/Users/alexi/.local/share/ipss/aplicacion_diccionario_intertrimestral_20261023_125556/ORIGENES_T1.json')
    original=Path('/Users/alexi/Downloads/RETENCION_IPSS_FINAL_con_Foto_Intertrimestral_20261022_AUDITADO_ENTREGA.xlsx')
    prom=Path('/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_12166.xlsx')
    bm=json.loads((R/'outputs/retencion_intermedia_t2_t3/20261021_115156_llaves_20_codcli/MANIFEST_DIAGNOSTICO.json').read_text())
    bref=next(x for x in bm['fuentes'] if 'IPSS_Procesamiento.xlsx' in x['ruta']);bridge=Path(bref['ruta'])
    assert sha(bridge)==bref['sha256']
    assert sha(originpath)=='9ba34a21ce89170f6343e068ba38ac1a546a927c1bd5ffcab9f4fd6ef48288a0'
    assert sha(original)=='20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8'
    assert sha(prom)=='aad1a71d166bdfb839049a5db534a21d050df1ed4e9251ad9982245896429015'
    origins=[x for x in json.loads(originpath.read_text()) if x['anio']==2025];assert len(origins)==1150

    def s(v):
     if v is None:return ''
     if isinstance(v,float) and v.is_integer():return str(int(v))
     return str(v).strip().upper().replace('.','').replace(' ','')
    def nrut(v):
     t=s(v)
     if not re.fullmatch(r'\d+-[0-9K]',t):return None
     b,d=t.split('-');return str(int(b))+'-'+d

    def rutpair(b,d):return nrut(s(b)+'-'+s(d))
    w=openpyxl.load_workbook(original,read_only=True,data_only=True)
    r9={i:r for i,r in enumerate(w['9_Base_Analitica_Detalle'].iter_rows(values_only=True),1)}
    r20={i:r for i,r in enumerate(w['20_Foto_BBDD'].iter_rows(values_only=True),1)}
    foto_t3=w['18_Foto_Intertrimestral']['G8'].value;assert foto_t3==1023
    s26=w['26_Avance_Curricular_Resumen'];notes={a:s26[a].value for a in ['A2','A1848','A1849','B1849','C1849','D1849','E1849','G1853','H1853']};old={c:s26[f'{c}1853'].value for c in ['C','D','E','F']};w.close()
    pp={s(o['documento_original']) for o in origins if o['tipo_doc']=='P'};assert len(pp)==20
    bridges=collections.defaultdict(list)
    b=openpyxl.load_workbook(bridge,read_only=True,data_only=True)
    for n,r in enumerate(b['2025_carga'].iter_rows(min_row=2,values_only=True),2):
     if len(r)<74:continue
     if s(r[72]) in pp:
      alias=nrut(r[9]);assert alias and rutpair(r[72],r[73])==alias
      bridges[s(r[72])].append({'alias':alias,'dv':s(r[73])})
    b.close()
    people=[];used=set();qa=collections.Counter()
    for o in origins:
     row20=r20[o['fila_foto']];rn=int(re.search(r'fila (\d+)',o['fuente']).group(1));z=r9[rn];typ=s(o['tipo_doc'])
     assert row20[1]==2025 and row20[2]==1 and z[1]==2025 and s(z[6])==typ
     assert o['fila_foto'] not in used;used.add(o['fila_foto'])
     if typ=='R':
      identity=rutpair(z[7],z[8]);assert identity and nrut(row20[4])==identity==o['persona'];qa['R_identidad_MU_y_hoja20_DV_exacto']+=1
     elif typ=='P':
      assert s(z[7])==s(o['documento_original'])==s(row20[4])==s(row20[8]);br=bridges[s(z[7])]
      aliases={x['alias'] for x in br};assert len(aliases)==1
      identity=next(iter(aliases));assert identity==o['persona'];qa['P_puente_explicito_N_DOC_J_BV']+=1
      qa['P_DV_ausente_en_MU']+=int(not s(z[8]))
     else:raise AssertionError('Tipo documental inesperado')
     people.append({'identity':identity,'type':typ,'code':o['codigo_origen'],'career':o['carrera']})
    assert len({x['identity'] for x in people})==1150
    assert len({x['identity'].split('-')[0] for x in people})==1150
    assert not ({x['identity'] for x in people if x['type']=='R'}&{x['identity'] for x in people if x['type']=='P'})
    identities={x['identity'] for x in people};bodies={x['identity'].split('-')[0] for x in people}
    pres=set();vig=set();bodypres=set();bodyvig=set();raw=collections.Counter();bad_dv=collections.Counter();matched_codcli=collections.defaultdict(set)
    p=openpyxl.load_workbook(prom,read_only=True,data_only=True);sheet=p['Hoja1'];h=list(next(sheet.iter_rows(max_row=1,values_only=True)))
    for n,r in enumerate(sheet.iter_rows(min_row=2,values_only=True),2):
     d=dict(zip(h,r))
     if d['ANO']!=2025 or d['PERIODO']!=3:continue
     full=rutpair(d['RUT'],d['DIG']);body=s(d['RUT']);isvig=s(d['ESTADO_ACADEMICO'])=='VIGENTE'
     if body in bodies:
      bodypres.add(body)
      if isvig:bodyvig.add(body)
      if full not in identities:bad_dv['filas_con_cuerpo_cohorte_y_DV_no_coincidente']+=1
     if full not in identities:continue
     pres.add(full);raw['filas_coincidentes_por_documento_y_DV']+=1;matched_codcli[full].add(s(d['CODCLI']))
     if isvig:vig.add(full);raw['filas_VIGENTE_coincidentes_por_documento_y_DV']+=1
    p.close()
    groups=collections.defaultdict(list)
    for x in people:groups[(x['code'],x['career'])].append(x)
    rows=[]
    for (code,career),members in sorted(groups.items()):
     rr={'codigo':code,'nombre':career,'denom':len(members),'denR':sum(x['type']=='R' for x in members),'denP':sum(x['type']=='P' for x in members)}
     for typ in ('R','P'):
      rr['obs'+typ]=sum(x['identity'] in pres and x['type']==typ for x in members)
      rr['vig'+typ]=sum(x['identity'] in vig and x['type']==typ for x in members)
     rr['observados']=rr['obsR']+rr['obsP'];rr['numerador_propuesto']=rr['vigR']+rr['vigP'];rr['tasa_propuesta']=rr['numerador_propuesto']/rr['denom'];rows.append(rr)
    tot={k:sum(x[k] for x in rows) for k in ['denom','denR','denP','obsR','obsP','vigR','vigP','observados','numerador_propuesto']}
    assert len(rows)==36 and tot['denom']==1150
    assert tot['numerador_propuesto']==546 and tot['observados']==596
    assert tot['vigR']==535 and tot['vigP']==11 and tot['denR']==1130 and tot['denP']==20
    assert tot['obsR']==582 and tot['obsP']==14
    res={'clasificacion':'B valores y fuentes; C reproducción nueva por identidad completa; E mecanismo original de535 no localizado','alcance':'Reproducción de auxiliar PROMEDIOS T3-2025, sin aplicar cambios. Foto T3=1023 se conserva separada e intacta.','metodo':'Cohorte2025 grupo1 completa,1150 personas. R: documento y DV de MU/hoja20; P: puente explícito IPSS_Procesamiento J/BU/BV. Intersección exacta RUT+DIG con PROMEDIOS ANO=2025 PERIODO=3 y al menos una fila ESTADO_ACADEMICO=VIGENTE. Personas únicas; sin filtro de SIES, CODCLI, plan, carrera ni versión de destino.','totales':tot,'tasa_actual':old['E']/old['C'],'tasa_propuesta':tot['numerador_propuesto']/tot['denom'],'diferencia_numerador':tot['numerador_propuesto']-old['E'],'por_codigo_T1':rows,'QA':dict(qa)|dict(raw)|dict(bad_dv)|{'personas_unicas_identidad_completa':len(identities),'personas_unicas_cuerpo':len(bodies),'colisiones_R_P':0,'duplicados_denominador':0,'filas_actividad_se_reducen_a_personas':True,'resultado_por_cuerpo_igual_a_identidad_completa':len(bodypres)==len(pres) and len(bodyvig)==len(vig),'personas_con_mas_de_un_CODCLI_en_registros_coincidentes':sum(len(x)>1 for x in matched_codcli.values()),'sin_nuevos_integrantes':True,'sin_cambios_fuentes':sha(original)=='20fe8438cba065f15240f1e67c066e5feb8bac8841f182864e40ca7bc9512aa8' and sha(prom)=='aad1a71d166bdfb839049a5db534a21d050df1ed4e9251ad9982245896429015'},'notas_observadas_hoja26':notes,'linaje_heredado':{'valores_fila1853':old,'constructor_temporal':str(constructor_original)+':43-46','observacion':'El constructor inserta535 literal como precomputed exact values; no contiene el algoritmo que calculó ese535. Hoja26 no contiene nómina individual de535. No se prueba causalidad de filtroP antiguo.','restriccion_tipo_documento_en_declaracion_disponible':False},'manejo_P':'MU/hoja9 no informa DV para los20 documentosP. No se inventa: se utiliza correspondencia explícita entre N_DOC en BU y RUT en J con DV en BV; en PROMEDIOS deben coincidir RUT y DIG completos. No se trata el vacíoMU como DV real.','fuentes':{str(x):{'sha256':sha(x)} for x in [original,originpath,bridge,prom,constructor_original]},'referencias':{'cohorte':'20_Foto_BBDD grupo1,año2025; origen original de cada integrante validado por fila9; referencias personales permanecen en memoria','puente':'IPSS_Procesamiento.xlsx/2025_carga: J RUT,BU N_DOC,BV DV','actividad':'PROMEDIOSDEALUMNOS_12166.xlsx/Hoja1: RUT,DIG,CODCLI,ANO,PERIODO,ESTADO_ACADEMICO','proxy_heredado':'26_Avance_Curricular_Resumen!A1848:H1855; E1853=535 literal; F1853=E1853/C1853'},'limites':['Se confirma nuevo conteo exacto uniforme; no se reconstruye el mecanismo original de535.','VIGENTE del extractoPROMEDIOS no acredita vigencia histórica a fecha exacta de cierre.','Identidad y presencia no autorizan cambiar denominadores ni atribuir versiónSIES.']}
    res['foto_T3_conservada_sin_modificar']=foto_t3
    res['derivacion_script']={'base':'scripts/recalcular_proxy_t2_2025_v14_20261023.py','sha256_base':sha(R/'scripts/recalcular_proxy_t2_2025_v14_20261023.py'),'cambio':'PERIODO3 y fila1853; mismo algoritmo exacto de identidad y puenteP.'}
    txt=json.dumps(res,ensure_ascii=False,indent=2)
    assert not any(x['identity'] in txt or x['identity'].split('-')[0] in txt for x in people)
    O=args.output_dir/'RESULTADO_AGREGADO.json';O.parent.mkdir(parents=True,exist_ok=True);O.write_text(txt+'\n')
    print(json.dumps({'totales':tot,'QA':res['QA'],'tasa_actual':res['tasa_actual'],'tasa_propuesta':res['tasa_propuesta'],'sha256':sha(O)},ensure_ascii=False,indent=2))


if __name__=='__main__':
    main()
