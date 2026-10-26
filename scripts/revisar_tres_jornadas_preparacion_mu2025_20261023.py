"""Contrasta tres pendientes con preparación MU2025 y contactos fechados; no adjudica."""
from pathlib import Path
import json,sys,datetime,re,subprocess,collections
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,rut_normalizado as normalizar_rut
read=lambda p:json.loads(Path(p).read_text())
now=datetime.datetime.now(ZoneInfo('America/Santiago'));O=R/'outputs'/('revision_tres_jornadas_'+now.strftime('%Y%m%d_%H%M%S'));P=Path.home()/'.local/share/ipss'/O.name
assert subprocess.run(['git','-C',str(P.parent),'rev-parse','--show-toplevel'],capture_output=True).returncode!=0
O.mkdir();P.mkdir(mode=0o700)
def dump(p,x):
 p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str)+'\n')
 if p.is_relative_to(P):p.chmod(0o600)
prior=R/'outputs/cierre_dos_jornadas_20261023_145844';m=read(prior/'MANIFEST.json');ts=read(Path(m['privado'])/'TRES_PENDIENTES.json');persons={x['persona']:x['caso'] for x in ts};bodies={p.split('-')[0]:c for p,c in persons.items()}
assert len(ts)==3 and len(persons)==2
inv=read(R/'outputs/retencion_intermedia_t2_t3/20261021_093412_git_onedrive_mu2025/RESPALDOS_XLSX_INSPECCION.json')
selected=[x for x in inv if Path(x['ruta']).name in ['IPSS_Procesamiento.xlsx','Formati_SIES_2025.xlsx','IPSS_2025.xlsx','Descarga de lo cargado_08-05-2025_22-26-22_Matrícula_Pregrado_2025.xlsx']]
fields={'CODCARPR','CODCARR','CODCLI','CODIGOCARRERA','CODCARRERA','ANOMATRICULA','PERIODOMATRICULA','ANOINGRESO','PERIODOINGRESO','FECHAMATRICULA','JORNADA','REGIMEN','PLAN_DE_ESTUDIO','COD_IES','COD_SED','COD_CAR','MODALIDAD','JOR','VERSION','FOR_ING_ACT','ANIO_ING_ACT','SEM_ING_ACT','ANIO_ING_ORI','SEM_ING_ORI','FECHA_MATRICULA','REINCORPORACION','VIG','ESTADOACADEMICO','NIVEL','CONTINUIDAD','ARREGLO','OBS','CARGA','FECH_CAR'}
records=[];inventory=[];hashes={str(prior/'MANIFEST.json'):sha256(prior/'MANIFEST.json')};allformulas=[]
for x in selected:
 p=Path(x['ruta']);assert sha256(p)==x['sha256'];hashes[str(p)]=x['sha256'];wv=load_workbook(p,read_only=True,data_only=True);wf=load_workbook(p,read_only=True,data_only=False)
 for ws in wv:
  if ws.max_row<90:continue
  header=None
  for rn,row in enumerate(ws.iter_rows(max_row=15,values_only=True),1):
   h=[str(v).strip().upper() if v is not None else '' for v in row]
   if 'N_DOC' in h or 'RUT' in h:header=h;start=rn;break
  if header is None:inventory.append({'fuente':str(p),'hoja':ws.title,'estado':'SIN_ENCABEZADO_IDENTIDAD_EN_15_FILAS'});continue
  inds=[i for i,h in enumerate(header) if h in ['N_DOC','RUT']];dv=next((i for i,h in enumerate(header) if h=='DV'),None);matches=[]
  for rn,row in enumerate(ws.iter_rows(min_row=start+1,values_only=True),start+1):
   found=set()
   for i in inds:
    if i>=len(row):continue
    val=row[i]
    if header[i]=='N_DOC' and dv is not None:
     val=str(int(val) if isinstance(val,(int,float)) else val)+'-'+str(row[dv])
    try:person=normalizar_rut(val)
    except ValueError:continue
    if person in persons:found.add(person)
   if not found:continue
   assert len(found)==1
   person=next(iter(found));vals={get_column_letter(i+1)+':'+h:row[i] for i,h in enumerate(header) if h in fields and i<len(row)}
   rec={'caso':persons[person],'persona':person,'fuente':str(p),'sha256':x['sha256'],'hoja':ws.title,'fila':rn,'valores':vals,'formulas':{}};records.append(rec);matches.append(rec)
  if matches:
   targetrows={z['fila']:z for z in matches}
   for rn,row in enumerate(wf[ws.title].iter_rows(min_row=min(targetrows),max_row=max(targetrows),values_only=True),min(targetrows)):
    if rn not in targetrows:continue
    for i,h in enumerate(header):
     if h in fields and i<len(row) and isinstance(row[i],str) and row[i].startswith('='):targetrows[rn]['formulas'][get_column_letter(i+1)+':'+h]=row[i]
  inventory.append({'fuente':str(p),'hoja':ws.title,'encabezado':start,'filas_foco':len(matches)})
 wf.close();wv.close();print(p.name,'leído',flush=True)
# Contactos: identidad exacta por cuerpo del RUT, sin reconstruir DV; solo notas académicas pertinentes.
pm=read(R/'outputs/secciones_cinco_jornadas_20261023_144650/MANIFEST.json');sg=next(Path(x) for x in pm['fuentes_sha256'] if Path(x).name=='GESTIONSGE_1703.xlsx');hashes[str(sg)]=pm['fuentes_sha256'][str(sg)];assert sha256(sg)==hashes[str(sg)]
w=load_workbook(sg,read_only=True,data_only=True);s=w['Hoja1'];h=next(s.iter_rows(values_only=True));contacts=[];rawcounts=collections.Counter()
for rn,row in enumerate(s.iter_rows(min_row=2,values_only=True),2):
 val=row[0];body=str(int(val) if isinstance(val,(int,float)) else val).strip().replace('.','');case=bodies.get(body)
 if not case:continue
 rawcounts[case]+=1
 obs=str(row[3] or '');date=row[7]
 if str(row[4])!='IINF' or not any(t in obs.lower() for t in ['jornada','ornada','diurn','vespert','plan','inscri','ramo','asignatur','modalidad','curso']):continue
 if any(t in obs.lower() for t in ['cobranza','moros','arancel','pagar','deuda']):continue
 contacts.append({'caso':case,'fuente':str(sg),'hoja':'Hoja1','fila':rn,'fecha':date,'CODCARR':row[4],'observacion':obs,'metodo':'Cuerpo RUT original exacto + CODCARR; DV no incluido en fuente'})
w.close()
dump(P/'EVIDENCIA.json',{'objetivos':ts,'preparacion_mu2025':records,'contactos':contacts})
# Publicar solo celdas académicas y etiquetas de caso, nunca identificadores ni observaciones libres.
public=[]
for z in records:
 vals={k:v for k,v in z['valores'].items() if k.split(':',1)[1] not in ['CODCLI','OBS','CARGA','ARREGLO']}
 public.append({k:z[k] for k in ['caso','fuente','sha256','hoja','fila']}|{'valores':vals,'formulas':z['formulas']})
dump(O/'PREPARACION_MU2025.json',public);dump(O/'INVENTARIO.json',inventory)
dump(O/'CONFIG.json',{'privado':str(P),'fuentes_sha256':hashes,'inicio':now.isoformat(),'diccionario_activo':read(R/'data/diccionario_sies_seguimiento/activo.json')})
print(json.dumps({'output':str(O),'filas_preparacion':len(records),'contactos_academicos':len(contacts),'contactos_por_caso':dict(collections.Counter(x['caso'] for x in contacts))},ensure_ascii=False))
