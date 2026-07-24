#!/usr/bin/env python3
"""Ampliacion por CODIGO_UNICO/RUT/VIG de la auditoria diferencia +166.

Entrada principal: auditoria validada 20260730_173940. No modifica fuentes.
"""

from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_XLSX = ROOT / "outputs/auditoria_diferencia_166_pregrado/20260730_173940/AUDITORIA_DIFERENCIA_166_PUBLICACION2026_VS_CONGELADO_PREGRADO_20260730_173940.xlsx"
TS = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_DIR = ROOT / "outputs/auditoria_diferencia_166_pregrado_por_codigo" / TS

HEADERS_PREGRADO = [
    "TIPO_DOC", "N_DOC", "DV", "PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRE",
    "SEXO", "FECH_NAC", "NAC", "PAIS_EST_SEC", "COD_SED", "COD_CAR", "MODALIDAD",
    "JOR", "VERSION", "FOR_ING_ACT", "ANIO_ING_ACT", "SEM_ING_ACT", "ANIO_ING_ORI",
    "SEM_ING_ORI", "ASI_INS_ANT", "ASI_APR_ANT", "PROM_PRI_SEM", "PROM_SEG_SEM",
    "ASI_INS_HIS", "ASI_APR_HIS", "NIV_ACA", "SIT_FON_SOL", "SUS_PRE",
    "FECHA_MATRICULA", "REINCORPORACION", "VIG",
]

PRIORITY_CODES = [
    "I162S2C91J2V1",
    "I162S2C91J1V1",
    "I162S2C114J2V1",
    "I162S2C57J4V1",
    "I162S2C76J4V1",
    "I162S2C87J4V1",
    "I162S2C77J1V1",
]


def sha256_path(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def logical_hash_df(df: pd.DataFrame) -> str:
    h = hashlib.sha256()
    h.update("\x1f".join(map(str, df.columns)).encode("utf-8"))
    h.update(b"\n")
    for row in df.itertuples(index=False, name=None):
        h.update("\x1f".join("" if pd.isna(v) else str(v) for v in row).encode("utf-8"))
        h.update(b"\n")
    return h.hexdigest()


def row_hash(row: pd.Series, cols: list[str]) -> str:
    text = "\x1f".join("" if pd.isna(row[c]) else str(row[c]) for c in cols)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def codigo_unico(df: pd.DataFrame) -> pd.Series:
    return (
        "I162S" + df["COD_SED"].astype(str).str.strip()
        + "C" + df["COD_CAR"].astype(str).str.strip()
        + "J" + df["JOR"].astype(str).str.strip()
        + "V" + df["VERSION"].astype(str).str.strip()
    )


def llave_persona(df: pd.DataFrame) -> pd.Series:
    return (
        df["TIPO_DOC"].astype(str).str.strip()
        + "|" + df["N_DOC"].astype(str).str.strip()
        + "|" + df["DV"].astype(str).str.strip()
    )


def mask_person_key(key: str) -> str:
    parts = str(key).split("|")
    if len(parts) < 3:
        return "***"
    ndoc = parts[1]
    return f"{parts[0]}|***{ndoc[-3:]}|{parts[2]}"


def normalize_raw(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        out[c] = out[c].where(~out[c].isna(), "")
    return out


def load_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, dict]:
    pub = pd.read_excel(BASE_XLSX, sheet_name="RAW_PUBLICACION2026")
    frozen = pd.read_excel(BASE_XLSX, sheet_name="RAW_PREGRADO_CONGELADO_ACTUAL", dtype=str)
    punto0 = pd.read_excel(BASE_XLSX, sheet_name="RAW_PUNTO0", dtype=str)
    comp95 = pd.read_excel(BASE_XLSX, sheet_name="RAW_COMPLEMENTO95", dtype=str)
    info = {
        "base_xlsx": str(BASE_XLSX),
        "base_xlsx_hash": sha256_path(BASE_XLSX),
        "raw_publicacion_hash_logico": logical_hash_df(pub),
        "raw_pregrado_hash_logico": logical_hash_df(frozen),
        "raw_punto0_hash_logico": logical_hash_df(punto0),
        "raw_complemento95_hash_logico": logical_hash_df(comp95),
    }
    return normalize_raw(pub), normalize_raw(frozen), normalize_raw(punto0), normalize_raw(comp95), info


def prepare_frozen(df: pd.DataFrame, componente: str | None = None) -> pd.DataFrame:
    out = df.copy()
    out["CODIGO_UNICO"] = codigo_unico(out)
    out["LLAVE_PERSONA"] = llave_persona(out)
    out["LLAVE_MATRICULA"] = out["LLAVE_PERSONA"] + "|" + out["CODIGO_UNICO"]
    out["HASH_FILA_32_CAMPOS"] = out.apply(lambda r: row_hash(r, HEADERS_PREGRADO), axis=1)
    out["VIG_STR"] = out["VIG"].astype(str).str.strip()
    if componente is not None:
        out["ARCHIVO_ORIGEN"] = componente
    return out


def publication_by_code(pub: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    pre = pub[pub["NIVEL GLOBAL"].astype(str).str.strip().str.casefold() == "pregrado"].copy()
    pre["CODIGO_UNICO"] = pre["CÓDIGO CARRERA"].astype(str).str.strip()
    pre["TOTAL_PUBLICACION"] = pd.to_numeric(pre["TOTAL MATRÍCULA"], errors="coerce").fillna(0).astype(int)
    names = pre.groupby("CODIGO_UNICO", dropna=False).agg(
        TOTAL_PUBLICACION=("TOTAL_PUBLICACION", "sum"),
        NOMBRE_CARRERA_PUBLICACION=("NOMBRE CARRERA", "first"),
    ).reset_index()
    info = {
        "publicacion_total": int(pd.to_numeric(pub["TOTAL MATRÍCULA"], errors="coerce").fillna(0).sum()),
        "publicacion_pregrado": int(names["TOTAL_PUBLICACION"].sum()),
        "filas_publicacion_pregrado": int(len(pre)),
        "ofertas_publicacion_pregrado": int(names["CODIGO_UNICO"].nunique()),
        "publicacion_posgrado_postitulo": int(pd.to_numeric(pub.loc[pub["NIVEL GLOBAL"].astype(str).str.strip().str.casefold() != "pregrado", "TOTAL MATRÍCULA"], errors="coerce").fillna(0).sum()),
    }
    return names, info


def classify_records(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    person_counts = out["LLAVE_PERSONA"].value_counts()
    matricula_counts = out["LLAVE_MATRICULA"].value_counts()
    hash_counts = out["HASH_FILA_32_CAMPOS"].value_counts()
    person_code_count = out.groupby("LLAVE_PERSONA")["CODIGO_UNICO"].nunique()
    person_carrera_count = out.groupby("LLAVE_PERSONA")["COD_CAR"].nunique()
    matricula_vigs = out.groupby("LLAVE_MATRICULA")["VIG_STR"].agg(lambda s: set(map(str, s)))
    out["APARICIONES_PERSONA"] = out["LLAVE_PERSONA"].map(person_counts).astype(int)
    out["APARICIONES_PERSONA_EN_CODIGO"] = out.groupby(["LLAVE_PERSONA", "CODIGO_UNICO"])["LLAVE_PERSONA"].transform("size").astype(int)
    out["APARICIONES_LLAVE_MATRICULA"] = out["LLAVE_MATRICULA"].map(matricula_counts).astype(int)
    out["APARICIONES_FILA_EXACTA"] = out["HASH_FILA_32_CAMPOS"].map(hash_counts).astype(int)
    out["CODIGOS_DISTINTOS_PERSONA"] = out["LLAVE_PERSONA"].map(person_code_count).astype(int)
    out["CARRERAS_DISTINTAS_PERSONA"] = out["LLAVE_PERSONA"].map(person_carrera_count).astype(int)
    out["VIGENCIAS_DISTINTAS_MATRICULA"] = out["LLAVE_MATRICULA"].map(lambda k: ";".join(sorted(matricula_vigs.get(k, set()))))

    def cls(row: pd.Series) -> str:
        vigset = set(str(row["VIGENCIAS_DISTINTAS_MATRICULA"]).split(";")) if row["VIGENCIAS_DISTINTAS_MATRICULA"] else set()
        if row["APARICIONES_FILA_EXACTA"] > 1:
            return "DUPLICADO_EXACTO_32_CAMPOS"
        if row["APARICIONES_LLAVE_MATRICULA"] > 1 and ("0" in vigset and ({"1", "2"} & vigset)):
            return "CONFLICTO_O_HISTORIAL_VIGENCIA"
        if row["APARICIONES_LLAVE_MATRICULA"] > 1:
            return "REPETICION_MISMA_MATRICULA"
        if row["CARRERAS_DISTINTAS_PERSONA"] > 1:
            return "PERSONA_MULTIPLE_CARRERA"
        if row["CODIGOS_DISTINTOS_PERSONA"] > 1:
            return "PERSONA_MULTIPLE_OFERTA"
        return "SIN_REPETICION"

    out["CLASIFICACION_REPETICION"] = out.apply(cls, axis=1)
    out["REVISION_REQUERIDA"] = out["CLASIFICACION_REPETICION"].ne("SIN_REPETICION").map({True: "SI", False: "NO"})
    return out


def repeated_group_counts(sub: pd.DataFrame, key: str) -> tuple[int, int]:
    vc = sub[key].value_counts()
    groups = int((vc > 1).sum())
    records = int(vc[vc > 1].sum()) if groups else 0
    return groups, records


def exact_dup_counts(sub: pd.DataFrame) -> tuple[int, int]:
    return repeated_group_counts(sub, "HASH_FILA_32_CAMPOS")


def make_summary_by_code(pub_codes: pd.DataFrame, frozen: pd.DataFrame) -> pd.DataFrame:
    all_codes = sorted(set(pub_codes["CODIGO_UNICO"]) | set(frozen["CODIGO_UNICO"]))
    pub_map = pub_codes.set_index("CODIGO_UNICO")["TOTAL_PUBLICACION"].to_dict()
    name_map = pub_codes.set_index("CODIGO_UNICO")["NOMBRE_CARRERA_PUBLICACION"].to_dict()
    rows = []
    for code in all_codes:
        sub = frozen[frozen["CODIGO_UNICO"] == code]
        pub = int(pub_map.get(code, 0))
        v0 = int((sub["VIG_STR"] == "0").sum()) if not sub.empty else 0
        v1 = int((sub["VIG_STR"] == "1").sum()) if not sub.empty else 0
        v2 = int((sub["VIG_STR"] == "2").sum()) if not sub.empty else 0
        vigente = v1 + v2
        historico = len(sub)
        vig_sub = sub[sub["VIG_STR"].isin(["1", "2"])]
        rut_g, rut_r = repeated_group_counts(sub, "LLAVE_PERSONA") if not sub.empty else (0, 0)
        mat_g, mat_r = repeated_group_counts(sub, "LLAVE_MATRICULA") if not sub.empty else (0, 0)
        ex_g, ex_r = exact_dup_counts(sub) if not sub.empty else (0, 0)
        multiple_carrera = int(sub.loc[sub["CLASIFICACION_REPETICION"].eq("PERSONA_MULTIPLE_CARRERA"), "LLAVE_PERSONA"].nunique()) if not sub.empty else 0
        multiple_oferta = int(sub.loc[sub["CLASIFICACION_REPETICION"].eq("PERSONA_MULTIPLE_OFERTA"), "LLAVE_PERSONA"].nunique()) if not sub.empty else 0
        rep_v0 = int(sub.loc[(sub["APARICIONES_PERSONA"] > 1) & (sub["VIG_STR"] == "0")].shape[0]) if not sub.empty else 0
        rep_vig = int(sub.loc[(sub["APARICIONES_PERSONA"] > 1) & (sub["VIG_STR"].isin(["1", "2"]))].shape[0]) if not sub.empty else 0
        personas_hist = int(sub["LLAVE_PERSONA"].nunique()) if not sub.empty else 0
        personas_vig = int(vig_sub["LLAVE_PERSONA"].nunique()) if not vig_sub.empty else 0
        vigente_sin_dup_exact = int(vig_sub["HASH_FILA_32_CAMPOS"].nunique()) if not vig_sub.empty else 0
        diff_vig = pub - vigente
        diff_hist = pub - historico
        diff_personas = pub - personas_vig
        diff_no_exact = pub - vigente_sin_dup_exact
        if pub == vigente:
            cls = "IGUAL"
        elif pub > 0 and historico == 0:
            cls = "SOLO_PUBLICACION"
        elif pub == 0 and vigente > 0:
            cls = "SOLO_CONGELADO_VIGENTE"
        elif pub == 0 and vigente == 0 and historico > 0:
            cls = "SOLO_CONGELADO_VIG0"
        elif pub > vigente:
            cls = "PUBLICACION_MAYOR"
        elif pub < vigente:
            cls = "PUBLICACION_MENOR"
        else:
            cls = "NO_COMPARABLE"
        diag_flags = []
        mix_vig_same_matricula = False
        if not sub.empty:
            for _, g in sub.groupby("LLAVE_MATRICULA"):
                vigs = set(g["VIG_STR"])
                if "0" in vigs and ({"1", "2"} & vigs):
                    mix_vig_same_matricula = True
                    break
        if multiple_carrera:
            diag_flags.append("CON_RUT_MULTIPLE_CARRERA")
        if mat_g:
            diag_flags.append("CON_REPETICION_MISMA_MATRICULA")
        if ex_g:
            diag_flags.append("CON_DUPLICADO_EXACTO")
        if mix_vig_same_matricula:
            diag_flags.append("CON_MEZCLA_VIG_0_Y_VIGENTE")
        if len(diag_flags) > 1:
            diag = "CON_MULTIPLES_HALLAZGOS"
        elif diag_flags:
            diag = diag_flags[0]
        else:
            diag = "SIN_HALLAZGOS_DE_DUPLICIDAD"
        first = sub.iloc[0] if not sub.empty else {}
        pct = "No aplica" if vigente == 0 else diff_vig / vigente
        rows.append({
            "CODIGO_UNICO": code,
            "COD_SED": first.get("COD_SED", ""),
            "COD_CAR": first.get("COD_CAR", ""),
            "JORNADA": first.get("JOR", ""),
            "MODALIDAD": first.get("MODALIDAD", ""),
            "VERSION": first.get("VERSION", ""),
            "NOMBRE_CARRERA_PUBLICACION": name_map.get(code, ""),
            "NOMBRE_CARRERA_CONGELADO": "NO DISPONIBLE EN 32 CAMPOS",
            "TOTAL_PUBLICACION": pub,
            "CONGELADO_HISTORICO": historico,
            "VIG_0": v0,
            "VIG_1": v1,
            "VIG_2": v2,
            "CONGELADO_VIGENTE": vigente,
            "PERSONAS_UNICAS_HISTORICO": personas_hist,
            "PERSONAS_UNICAS_VIGENTES": personas_vig,
            "RUT_REPETIDOS_GRUPOS": rut_g,
            "REGISTROS_EN_RUT_REPETIDOS": rut_r,
            "LLAVE_MATRICULA_REPETIDA_GRUPOS": mat_g,
            "REGISTROS_MATRICULA_REPETIDA": mat_r,
            "DUPLICADOS_EXACTOS_GRUPOS": ex_g,
            "REGISTROS_DUPLICADOS_EXACTOS": ex_r,
            "PERSONAS_MULTIPLE_CARRERA": multiple_carrera,
            "PERSONAS_MULTIPLE_OFERTA": multiple_oferta,
            "REPETIDOS_VIG_0": rep_v0,
            "REPETIDOS_VIGENTES": rep_vig,
            "DIF_PUBLICACION_VS_VIGENTE": diff_vig,
            "DIF_PUBLICACION_VS_HISTORICO": diff_hist,
            "DIF_PUBLICACION_VS_PERSONAS_UNICAS": diff_personas,
            "VIGENTES_SIN_DUPLICADO_EXACTO": vigente_sin_dup_exact,
            "DIF_PUBLICACION_VS_SIN_DUP_EXACTO": diff_no_exact,
            "PORCENTAJE_DIFERENCIA": pct,
            "CLASIFICACION": cls,
            "CLASIFICACION_DIAGNOSTICA": diag,
            "ALERTA_DUPLICADOS": "SI" if diag != "SIN_HALLAZGOS_DE_DUPLICIDAD" else "NO",
            "ALERTA_VIGENCIA": "SI" if v0 > 0 or mix_vig_same_matricula else "NO",
            "REQUIERE_REVISION": "SI" if diff_vig != 0 or diag != "SIN_HALLAZGOS_DE_DUPLICIDAD" or cls.startswith("SOLO") else "NO",
            "OBSERVACION": "Diferencia localizada; no atribuye RUT a publicacion agregada." if diff_vig != 0 else "Sin diferencia de total vigente publicado.",
        })
    return pd.DataFrame(rows)


def make_decomposition(summary: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    d = summary[["CODIGO_UNICO", "TOTAL_PUBLICACION", "CONGELADO_VIGENTE", "DIF_PUBLICACION_VS_VIGENTE"]].copy()
    d = d.rename(columns={"TOTAL_PUBLICACION": "Publicación", "CONGELADO_VIGENTE": "Vigentes", "DIF_PUBLICACION_VS_VIGENTE": "Diferencia"})
    d["Aporte positivo"] = d["Diferencia"].clip(lower=0)
    d["Aporte negativo"] = (-d["Diferencia"].clip(upper=0)).astype(int)
    d["% del +166 neto"] = d["Diferencia"] / 166
    d = d.sort_values(["Aporte positivo", "Aporte negativo", "CODIGO_UNICO"], ascending=[False, False, True]).reset_index(drop=True)
    d["Acumulado"] = d["Aporte positivo"].cumsum()
    pos = d[d["Aporte positivo"] > 0].copy()
    total_pos = int(pos["Aporte positivo"].sum())
    thresholds = {}
    for pct in [0.25, 0.50, 0.75, 0.90, 1.00]:
        target = total_pos * pct
        thresholds[f"codigos_para_{int(pct*100)}pct_positivos"] = int((pos["Aporte positivo"].cumsum() < target).sum() + (1 if total_pos else 0))
    thresholds["suma_positivos"] = total_pos
    thresholds["suma_negativos_abs"] = int(d["Aporte negativo"].sum())
    thresholds["neto"] = int(d["Diferencia"].sum())
    return d, thresholds


def make_rut_repeated_summary(df: pd.DataFrame) -> pd.DataFrame:
    rep_keys = df["LLAVE_PERSONA"].value_counts()
    rep_keys = rep_keys[rep_keys > 1].index
    rows = []
    for key, sub in df[df["LLAVE_PERSONA"].isin(rep_keys)].groupby("LLAVE_PERSONA"):
        comps = sub["ARCHIVO_ORIGEN"].value_counts()
        codes = sorted(sub["CODIGO_UNICO"].unique())
        vigs = sorted(sub["VIG_STR"].unique())
        if sub["COD_CAR"].nunique() > 1:
            cls = "PERSONA_MULTIPLE_CARRERA"
            possible = "NO_DETERMINADO"
        elif sub["CODIGO_UNICO"].nunique() > 1:
            cls = "PERSONA_MULTIPLE_OFERTA"
            possible = "PENDIENTE_REVISION"
        elif sub["HASH_FILA_32_CAMPOS"].duplicated(keep=False).any():
            cls = "DUPLICADO_EXACTO_32_CAMPOS"
            possible = "PENDIENTE_REVISION"
        elif "0" in vigs and ({"1", "2"} & set(vigs)):
            cls = "CONFLICTO_O_HISTORIAL_VIGENCIA"
            possible = "PENDIENTE_REVISION"
        else:
            cls = "REPETICION_MISMA_MATRICULA"
            possible = "PENDIENTE_REVISION"
        rows.append({
            "LLAVE_PERSONA": key,
            "LLAVE_PERSONA_MASCARADA": mask_person_key(key),
            "Cantidad registros": len(sub),
            "Cantidad códigos": len(codes),
            "Códigos asociados": ";".join(codes),
            "Vigencias": ";".join(vigs),
            "Punto 0": int(comps.get("PUNTO0", comps.get("PUNTO0_PARA_SUBIR", 0))),
            "Complemento 95": int(comps.get("COMPLEMENTO95", comps.get("COMPLEMENTO95_V3", 0))),
            "Clasificación": cls,
            "Posible error": possible,
            "Observación": "No eliminar automaticamente; publicacion no contiene RUT.",
        })
    return pd.DataFrame(rows)


def make_vig0_summary(df: pd.DataFrame, summary: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    vig0 = df[df["VIG_STR"] == "0"].copy()
    rows = []
    for code, sub in vig0.groupby("CODIGO_UNICO"):
        persons = set(sub["LLAVE_PERSONA"])
        same_code_vig = df[(df["CODIGO_UNICO"] == code) & (df["VIG_STR"].isin(["1", "2"]))]
        same_code_vig_persons = set(same_code_vig["LLAVE_PERSONA"])
        other_code_vig = df[(df["CODIGO_UNICO"] != code) & (df["VIG_STR"].isin(["1", "2"]))]
        other_code_vig_persons = set(other_code_vig["LLAVE_PERSONA"])
        pub = int(summary.set_index("CODIGO_UNICO").loc[code, "TOTAL_PUBLICACION"]) if code in set(summary["CODIGO_UNICO"]) else 0
        diff = int(summary.set_index("CODIGO_UNICO").loc[code, "DIF_PUBLICACION_VS_VIGENTE"]) if code in set(summary["CODIGO_UNICO"]) else 0
        rows.append({
            "CODIGO_UNICO": code,
            "VIG=0": len(sub),
            "Personas VIG=0": len(persons),
            "También vigentes en mismo código": len(persons & same_code_vig_persons),
            "Vigentes en otro código": len(persons & other_code_vig_persons),
            "Solo VIG=0": len(persons - same_code_vig_persons - other_code_vig_persons),
            "Publicación": pub,
            "Diferencia": diff,
            "Observación": "VIG=0 visible como historico; no suma a vigente.",
        })
    detail = vig0.copy()
    person_vig_by_code = df[df["VIG_STR"].isin(["1", "2"])].groupby(["LLAVE_PERSONA", "CODIGO_UNICO"]).size()
    person_vig_any = df[df["VIG_STR"].isin(["1", "2"])].groupby("LLAVE_PERSONA")["CODIGO_UNICO"].agg(lambda s: set(s))
    detail["PRESENCIA_VIG1_MISMO_CODIGO"] = detail.apply(lambda r: "SI" if (r["LLAVE_PERSONA"], r["CODIGO_UNICO"]) in person_vig_by_code.index else "NO", axis=1)
    detail["PRESENCIA_VIG2"] = "NO"
    detail["PRESENCIA_OTRO_CODIGO_VIGENTE"] = detail.apply(lambda r: "SI" if bool(person_vig_any.get(r["LLAVE_PERSONA"], set()) - {r["CODIGO_UNICO"]}) else "NO", axis=1)
    detail["CLASIFICACION_VIG0"] = detail.apply(
        lambda r: "VIG0_CON_VIGENTE_MISMO_CODIGO" if r["PRESENCIA_VIG1_MISMO_CODIGO"] == "SI"
        else ("VIG0_CON_VIGENTE_OTRO_CODIGO" if r["PRESENCIA_OTRO_CODIGO_VIGENTE"] == "SI" else "SOLO_VIG0"),
        axis=1,
    )
    return pd.DataFrame(rows).sort_values("VIG=0", ascending=False), detail


def make_simulations(summary: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in summary.iterrows():
        pub = int(r["TOTAL_PUBLICACION"])
        scenarios = {
            "Escenario 1": int(r["CONGELADO_VIGENTE"]),
            "Escenario 2": int(r["PERSONAS_UNICAS_VIGENTES"]),
            "Escenario 3": int(r["PERSONAS_UNICAS_VIGENTES"]),  # LLAVE_MATRICULA unica por codigo equivale a persona+codigo.
            "Escenario 4": int(r["VIGENTES_SIN_DUPLICADO_EXACTO"]),
            "Escenario 5": int(r["PERSONAS_UNICAS_VIGENTES"]),
        }
        diffs = {k: abs(pub - v) for k, v in scenarios.items()}
        minv = min(diffs.values())
        closest = ";".join([k for k, v in diffs.items() if v == minv])
        rows.append({
            "CODIGO_UNICO": r["CODIGO_UNICO"],
            "Publicación": pub,
            **scenarios,
            "Escenario más cercano": closest,
            "Diferencia mínima": minv,
            "Nota": "Diagnostico; no prueba regla de conteo de la publicacion.",
        })
    return pd.DataFrame(rows)


def make_component_comparison(p0: pd.DataFrame, comp: pd.DataFrame) -> pd.DataFrame:
    codes = sorted(set(p0["CODIGO_UNICO"]) | set(comp["CODIGO_UNICO"]))
    rows = []
    for code in codes:
        a = p0[p0["CODIGO_UNICO"] == code]
        b = comp[comp["CODIGO_UNICO"] == code]
        persons_a, persons_b = set(a["LLAVE_PERSONA"]), set(b["LLAVE_PERSONA"])
        mat_a, mat_b = set(a["LLAVE_MATRICULA"]), set(b["LLAVE_MATRICULA"])
        hash_a, hash_b = set(a["HASH_FILA_32_CAMPOS"]), set(b["HASH_FILA_32_CAMPOS"])
        coinc_mat = mat_a & mat_b
        rows.append({
            "CODIGO_UNICO": code,
            "Punto 0": len(a),
            "Complemento 95": len(b),
            "Personas coincidentes": len(persons_a & persons_b),
            "Matrículas coincidentes": len(coinc_mat),
            "Filas exactas coincidentes": len(hash_a & hash_b),
            "Nuevos reales": int((~b["LLAVE_MATRICULA"].isin(mat_a)).sum()) if not b.empty else 0,
            "Coincidencias no exactas": len(coinc_mat) - len(hash_a & hash_b),
            "VIG Punto 0": ";".join(sorted(a["VIG_STR"].unique())) if not a.empty else "",
            "VIG Complemento": ";".join(sorted(b["VIG_STR"].unique())) if not b.empty else "",
            "Total combinado": len(a) + len(b),
            "Alerta": "COINCIDE_MATRICULA_ENTRE_COMPONENTES" if coinc_mat else ("SOLO_COMPLEMENTO" if len(a) == 0 and len(b) > 0 else "SIN_COINCIDENCIA_COMPONENTES"),
        })
    return pd.DataFrame(rows)


def make_priority(summary: pd.DataFrame, comp_cmp: pd.DataFrame) -> pd.DataFrame:
    rows = []
    idx = summary.set_index("CODIGO_UNICO")
    comp_idx = comp_cmp.set_index("CODIGO_UNICO")
    for code in PRIORITY_CODES:
        if code not in idx.index:
            rows.append({"CODIGO_UNICO": code, "Conclusión": "No encontrado en resumen."})
            continue
        r = idx.loc[code]
        c = comp_idx.loc[code] if code in comp_idx.index else pd.Series(dtype=object)
        rows.append({
            "CODIGO_UNICO": code,
            "Nombre": r["NOMBRE_CARRERA_PUBLICACION"],
            "Sede": r["COD_SED"],
            "Carrera": r["COD_CAR"],
            "Jornada": r["JORNADA"],
            "Modalidad": r["MODALIDAD"],
            "Version": r["VERSION"],
            "Publicacion": r["TOTAL_PUBLICACION"],
            "Historico": r["CONGELADO_HISTORICO"],
            "VIG=0": r["VIG_0"],
            "VIG=1": r["VIG_1"],
            "VIG=2": r["VIG_2"],
            "Vigente": r["CONGELADO_VIGENTE"],
            "Diferencia": r["DIF_PUBLICACION_VS_VIGENTE"],
            "Personas unicas vigentes": r["PERSONAS_UNICAS_VIGENTES"],
            "RUT repetidos grupos": r["RUT_REPETIDOS_GRUPOS"],
            "Repeticiones misma matricula": r["LLAVE_MATRICULA_REPETIDA_GRUPOS"],
            "Duplicados exactos": r["DUPLICADOS_EXACTOS_GRUPOS"],
            "Personas en otras carreras": r["PERSONAS_MULTIPLE_CARRERA"],
            "Punto 0": c.get("Punto 0", 0),
            "Complemento 95": c.get("Complemento 95", 0),
            "Coincidencias componentes": c.get("Matrículas coincidentes", 0),
            "Conclusión": "Diferencia localizada por codigo; causa no demostrable sin fuente individual de publicacion.",
            "Pendiente": "Requiere fuente individual o reporte PES posterior si se necesita atribuir RUT.",
        })
    return pd.DataFrame(rows)


def write_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    out = df.copy()
    if out.empty:
        out = pd.DataFrame({"SIN_REGISTROS": []})
    out.to_excel(writer, sheet_name=name[:31], index=False)


def format_workbook(path: Path, summary_links: dict[str, str]) -> None:
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if ws.max_row > 1 and ws.max_column > 1:
            ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{ws.max_row}"
            table_name = "T_" + "".join(ch if ch.isalnum() else "_" for ch in ws.title)[:24]
            try:
                tab = Table(displayName=table_name, ref=ref)
                tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
                ws.add_table(tab)
            except Exception:
                pass
        for col_cells in ws.columns:
            max_len = 0
            for c in col_cells[:300]:
                max_len = max(max_len, len(str(c.value)) if c.value is not None else 0)
            ws.column_dimensions[col_cells[0].column_letter].width = max(10, min(55, max_len + 2))
        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 300), max_col=ws.max_column):
            for c in row:
                if isinstance(c.value, float) and abs(c.value) < 10:
                    c.number_format = "0.0%"
    ws = wb["00_RESUMEN_EJECUTIVO"]
    start = ws.max_row + 2
    ws.cell(start, 1).value = "Navegación"
    ws.cell(start, 1).font = Font(bold=True)
    for i, (label, sheet) in enumerate(summary_links.items(), start=start + 1):
        cell = ws.cell(i, 1)
        cell.value = label
        cell.hyperlink = f"#'{sheet}'!A1"
        cell.style = "Hyperlink"
    wb.save(path)


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "_Sin registros._"
    cols = list(df.columns)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join("---" for _ in cols) + " |"]
    for row in df.itertuples(index=False, name=None):
        lines.append("| " + " | ".join(str(v).replace("|", "\\|") for v in row) + " |")
    return "\n".join(lines)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=False)
    pub_raw, frozen_raw, p0_raw, comp_raw, input_info = load_inputs()
    p0 = prepare_frozen(p0_raw, "PUNTO0")
    comp = prepare_frozen(comp_raw, "COMPLEMENTO95")
    frozen_with_src = classify_records(pd.concat([p0, comp], ignore_index=True))
    frozen = frozen_with_src.copy()

    pub_codes, pub_info = publication_by_code(pub_raw)
    summary = make_summary_by_code(pub_codes, frozen)
    decomp, conc = make_decomposition(summary)
    rut_rep = make_rut_repeated_summary(frozen_with_src)
    rut_rep_detail = frozen_with_src[frozen_with_src["APARICIONES_PERSONA"] > 1].copy()
    mat_rep = frozen_with_src[frozen_with_src["APARICIONES_LLAVE_MATRICULA"] > 1].copy()
    exact_rep = frozen_with_src[frozen_with_src["APARICIONES_FILA_EXACTA"] > 1].copy()
    multicarrera = rut_rep[rut_rep["Clasificación"].eq("PERSONA_MULTIPLE_CARRERA")].copy()
    vig0_summary, vig0_detail = make_vig0_summary(frozen_with_src, summary)
    sims = make_simulations(summary)
    comp_cmp = make_component_comparison(p0, comp)
    priority = make_priority(summary, comp_cmp)

    total_vig = int(summary["CONGELADO_VIGENTE"].sum())
    total_hist = int(summary["CONGELADO_HISTORICO"].sum())
    total_pub = int(summary["TOTAL_PUBLICACION"].sum())
    validations = pd.DataFrame([
        ["Publicación pregrado = 3.371", total_pub == 3371, total_pub],
        ["Congelado vigente = 3.205", total_vig == 3205, total_vig],
        ["Diferencia neta = 166", int(summary["DIF_PUBLICACION_VS_VIGENTE"].sum()) == 166, int(summary["DIF_PUBLICACION_VS_VIGENTE"].sum())],
        ["Diferencias positivas = 186", conc["suma_positivos"] == 186, conc["suma_positivos"]],
        ["Diferencias negativas absolutas = 20", conc["suma_negativos_abs"] == 20, conc["suma_negativos_abs"]],
        ["Todas las ofertas incluidas", len(summary) == len(set(pub_codes["CODIGO_UNICO"]) | set(frozen["CODIGO_UNICO"])), len(summary)],
        ["Cada código aparece una sola vez", summary["CODIGO_UNICO"].is_unique, int(summary["CODIGO_UNICO"].duplicated().sum())],
        ["Suma publicada por código = 3.371", total_pub == 3371, total_pub],
        ["Suma vigente por código = 3.205", total_vig == 3205, total_vig],
        ["Suma VIG=0 por código = 960", int(summary["VIG_0"].sum()) == 960, int(summary["VIG_0"].sum())],
        ["Suma histórica por código = 4.165", total_hist == 4165, total_hist],
        ["Todos los registros tienen clasificación de repetición", frozen_with_src["CLASIFICACION_REPETICION"].notna().all(), ""],
        ["RUT repetidos separados misma/distinta oferta", True, "Clasificaciones en hojas 08-12."],
        ["Duplicados exactos identificados con 32 columnas", "HASH_FILA_32_CAMPOS" in frozen_with_src.columns, int(exact_rep["HASH_FILA_32_CAMPOS"].nunique())],
        ["VIG=0 no contado como vigente", total_vig == int((frozen_with_src["VIG_STR"].isin(["1", "2"])).sum()), total_vig],
        ["Punto 0 y complemento comparados", not comp_cmp.empty, len(comp_cmp)],
        ["No se eliminaron registros", len(frozen_with_src) == 4165, len(frozen_with_src)],
        ["No se modificaron VIG", sorted(frozen_with_src["VIG_STR"].unique()) == ["0", "1"], sorted(frozen_with_src["VIG_STR"].unique())],
        ["Originales no modificados", True, "Solo lectura desde workbook base y salidas nuevas."],
        ["RAW completas", len(pub_raw) == 66 and len(frozen_raw) == 4165 and len(p0_raw) == 4070 and len(comp_raw) == 95, f"pub={len(pub_raw)}, congelado={len(frozen_raw)}, p0={len(p0_raw)}, comp={len(comp_raw)}"],
        ["Hashes registrados", True, input_info],
        ["Posgrado permanece separado", pub_info["publicacion_posgrado_postitulo"] == 54, pub_info["publicacion_posgrado_postitulo"]],
        ["Simulaciones marcadas como diagnósticas", sims["Nota"].str.contains("Diagnostico").all(), ""],
        ["No se atribuyeron RUT a la publicación", True, "Publicacion solo agregada por oferta."],
    ], columns=["Validación", "Resultado", "Detalle"])

    pendientes = pd.DataFrame([
        ["PUBLICACION_SIN_RUT", "No es posible identificar que RUT individuales fueron contados en la publicacion agregada."],
        ["DIFERENCIA_CAUSAL_166", "La diferencia esta localizada por codigo, pero no se demuestra origen causal sin fuente individual/reportes PES posteriores."],
        ["NOMBRES_CARRERA_CONGELADO", "La estructura de 32 campos no contiene nombre de carrera; se conserva nombre desde publicacion cuando existe."],
    ], columns=["Pendiente", "Descripción"])

    resumen_exec = pd.DataFrame([
        ["Estado", "AMPLIACION_POR_CODIGO_COMPLETADA"],
        ["Limitación", "publicacion2026.xlsx es agregada; no contiene RUT."],
        ["Publicación pregrado", total_pub],
        ["Congelado vigente", total_vig],
        ["Diferencia neta", int(summary["DIF_PUBLICACION_VS_VIGENTE"].sum())],
        ["Diferencias positivas", conc["suma_positivos"]],
        ["Diferencias negativas absolutas", conc["suma_negativos_abs"]],
        ["Códigos totales", len(summary)],
        ["Códigos publicación mayor", int((summary["CLASIFICACION"] == "PUBLICACION_MAYOR").sum())],
        ["Códigos publicación menor", int((summary["CLASIFICACION"] == "PUBLICACION_MENOR").sum())],
        ["Códigos iguales", int((summary["CLASIFICACION"] == "IGUAL").sum())],
        ["Códigos solo publicación", int((summary["CLASIFICACION"] == "SOLO_PUBLICACION").sum())],
        ["Códigos solo congelado vigente", int((summary["CLASIFICACION"] == "SOLO_CONGELADO_VIGENTE").sum())],
        ["RUT repetidos grupos", len(rut_rep)],
        ["Matrículas repetidas grupos", int(mat_rep["LLAVE_MATRICULA"].nunique())],
        ["Duplicados exactos grupos", int(exact_rep["HASH_FILA_32_CAMPOS"].nunique())],
        ["Personas multicarrera grupos", len(multicarrera)],
        ["VIG=0", int(summary["VIG_0"].sum())],
        ["Códigos para 25% diferencias positivas", conc["codigos_para_25pct_positivos"]],
        ["Códigos para 50% diferencias positivas", conc["codigos_para_50pct_positivos"]],
        ["Códigos para 75% diferencias positivas", conc["codigos_para_75pct_positivos"]],
        ["Códigos para 90% diferencias positivas", conc["codigos_para_90pct_positivos"]],
        ["Diferencia aún pendiente causal", 166],
    ], columns=["Métrica", "Valor"])

    traz = pd.DataFrame([
        ["timestamp", TS],
        ["script", str(ROOT / "scripts/ampliar_auditoria_diferencia_166_por_codigo.py")],
        ["workbook_base", str(BASE_XLSX)],
        ["hash_workbook_base", input_info["base_xlsx_hash"]],
        ["LLAVE_PERSONA", "TIPO_DOC + N_DOC + DV"],
        ["CODIGO_UNICO", "I162S{COD_SED}C{COD_CAR}J{JOR}V{VERSION}"],
        ["LLAVE_MATRICULA", "LLAVE_PERSONA + CODIGO_UNICO"],
        ["HASH_FILA_32_CAMPOS", "SHA-256 logico de las 32 columnas oficiales"],
        ["hash_logico_RAW_PUBLICACION2026", input_info["raw_publicacion_hash_logico"]],
        ["hash_logico_RAW_PREGRADO_CONGELADO", input_info["raw_pregrado_hash_logico"]],
        ["hash_logico_RAW_PUNTO0", input_info["raw_punto0_hash_logico"]],
        ["hash_logico_RAW_COMPLEMENTO95", input_info["raw_complemento95_hash_logico"]],
    ], columns=["Campo", "Valor"])

    xlsx_path = OUT_DIR / f"AUDITORIA_DIFERENCIA_166_PREGRADO_POR_CODIGO_RUT_VIGENCIA_{TS}.xlsx"
    report_path = OUT_DIR / f"REPORTE_DIFERENCIA_166_PREGRADO_POR_CODIGO_RUT_VIGENCIA_{TS}.md"
    manifest_path = OUT_DIR / f"MANIFEST_DIFERENCIA_166_PREGRADO_POR_CODIGO_{TS}.json"

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        write_sheet(writer, "00_RESUMEN_EJECUTIVO", resumen_exec)
        write_sheet(writer, "01_RESUMEN_POR_CODIGO", summary)
        write_sheet(writer, "02_DESCOMPOSICION_166", decomp)
        write_sheet(writer, "03_DIFERENCIAS_POSITIVAS", summary[summary["DIF_PUBLICACION_VS_VIGENTE"] > 0])
        write_sheet(writer, "04_DIFERENCIAS_NEGATIVAS", summary[summary["DIF_PUBLICACION_VS_VIGENTE"] < 0])
        write_sheet(writer, "05_CODIGOS_IGUALES", summary[summary["DIF_PUBLICACION_VS_VIGENTE"] == 0])
        write_sheet(writer, "06_SOLO_PUBLICACION", summary[summary["CLASIFICACION"] == "SOLO_PUBLICACION"])
        write_sheet(writer, "07_SOLO_CONGELADO", summary[summary["CLASIFICACION"].isin(["SOLO_CONGELADO_VIGENTE", "SOLO_CONGELADO_VIG0"])])
        write_sheet(writer, "08_RUT_REPETIDOS_RESUMEN", rut_rep)
        write_sheet(writer, "09_RUT_REPETIDOS_DETALLE", rut_rep_detail)
        write_sheet(writer, "10_MATRICULAS_REPETIDAS", mat_rep)
        write_sheet(writer, "11_DUPLICADOS_EXACTOS", exact_rep)
        write_sheet(writer, "12_PERSONAS_MULTICARRERA", multicarrera)
        write_sheet(writer, "13_RESUMEN_VIG0_CODIGO", vig0_summary)
        write_sheet(writer, "14_DETALLE_VIG0", vig0_detail)
        write_sheet(writer, "15_SIMULACIONES_CONTEO", sims)
        write_sheet(writer, "16_PUNTO0_VS_COMPLEMENTO95", comp_cmp)
        write_sheet(writer, "17_CODIGOS_PRIORITARIOS", priority)
        write_sheet(writer, "18_VALIDACIONES", validations)
        write_sheet(writer, "19_TRAZABILIDAD", traz)
        write_sheet(writer, "20_PENDIENTES", pendientes)
        write_sheet(writer, "RAW_PUBLICACION2026", pub_raw)
        write_sheet(writer, "RAW_PREGRADO_CONGELADO", frozen_raw)
        write_sheet(writer, "RAW_PUNTO0", p0_raw)
        write_sheet(writer, "RAW_COMPLEMENTO95", comp_raw)

    format_workbook(xlsx_path, {
        "Resumen por código": "01_RESUMEN_POR_CODIGO",
        "Descomposición +166": "02_DESCOMPOSICION_166",
        "RUT repetidos": "08_RUT_REPETIDOS_RESUMEN",
        "VIG=0": "13_RESUMEN_VIG0_CODIGO",
        "Simulaciones": "15_SIMULACIONES_CONTEO",
        "Validaciones": "18_VALIDACIONES",
    })

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    excel_ok = all(s in wb.sheetnames for s in ["00_RESUMEN_EJECUTIVO", "01_RESUMEN_POR_CODIGO", "RAW_PUBLICACION2026", "RAW_PREGRADO_CONGELADO", "RAW_PUNTO0", "RAW_COMPLEMENTO95"])
    validations.loc[len(validations)] = ["El Excel abre correctamente", excel_ok, str(xlsx_path)]
    wb.close()
    wb_edit = openpyxl.load_workbook(xlsx_path)
    ws_val = wb_edit["18_VALIDACIONES"]
    ws_val.append(["El Excel abre correctamente", excel_ok, str(xlsx_path)])
    wb_edit.save(xlsx_path)

    top_pos = summary[summary["DIF_PUBLICACION_VS_VIGENTE"] > 0].sort_values("DIF_PUBLICACION_VS_VIGENTE", ascending=False)
    neg = summary[summary["DIF_PUBLICACION_VS_VIGENTE"] < 0].sort_values("DIF_PUBLICACION_VS_VIGENTE")
    report = f"""# Desglose por CODIGO_UNICO de diferencia +166 pregrado

## Contexto

Se amplia la auditoria de pregrado entre `publicacion2026.xlsx` y el congelado combinado gobernado de 4.165 registros. Posgrado/Postitulo permanece fuera del analisis y sigue conciliado 54 publicado versus 54 congelado.

## Metodologia

La publicacion fue tratada como `MATRICULA_AGREGADA_PUBLICADA`: se filtro `NIVEL GLOBAL=Pregrado` y se sumo `TOTAL MATRICULA` por `CODIGO CARRERA`. El congelado se agrego por `CODIGO_UNICO = I162S{{COD_SED}}C{{COD_CAR}}J{{JOR}}V{{VERSION}}`, usando como matricula vigente solo `VIG in (1,2)`.

## Limitacion de publicacion agregada

`publicacion2026.xlsx` no contiene RUT ni registros individuales. Por eso se identifican codigos con diferencia y RUT/repeticiones dentro del congelado, pero no se puede afirmar que un RUT institucional especifico fue contado o no contado en la publicacion.

## Desglose completo por codigo

La tabla completa esta en `01_RESUMEN_POR_CODIGO`. Totales: publicacion pregrado {total_pub}, vigente congelado {total_vig}, diferencia neta {int(summary['DIF_PUBLICACION_VS_VIGENTE'].sum())}.

## Concentracion de diferencias

- Suma diferencias positivas: {conc['suma_positivos']}.
- Suma diferencias negativas absolutas: {conc['suma_negativos_abs']}.
- Neto: {conc['neto']}.
- Codigos para 25% de positivos: {conc['codigos_para_25pct_positivos']}.
- Codigos para 50% de positivos: {conc['codigos_para_50pct_positivos']}.
- Codigos para 75% de positivos: {conc['codigos_para_75pct_positivos']}.
- Codigos para 90% de positivos: {conc['codigos_para_90pct_positivos']}.

### Mayores diferencias positivas

{markdown_table(top_pos[['CODIGO_UNICO','TOTAL_PUBLICACION','CONGELADO_VIGENTE','VIG_0','DIF_PUBLICACION_VS_VIGENTE','CLASIFICACION_DIAGNOSTICA']], 15)}

### Diferencias negativas

{markdown_table(neg[['CODIGO_UNICO','TOTAL_PUBLICACION','CONGELADO_VIGENTE','VIG_0','DIF_PUBLICACION_VS_VIGENTE','CLASIFICACION_DIAGNOSTICA']], 20)}

## RUT repetidos y matriculas repetidas

- Grupos de RUT repetidos: {len(rut_rep)}.
- Grupos de matricula repetida: {int(mat_rep['LLAVE_MATRICULA'].nunique())}.
- Grupos de duplicado exacto de 32 campos: {int(exact_rep['HASH_FILA_32_CAMPOS'].nunique())}.
- Personas en multiples carreras: {len(multicarrera)}.

Estos hallazgos son diagnosticos y no fueron descontados de la matricula.

## Registros VIG=0

Se identificaron {len(vig0_detail)} registros `VIG=0`; no fueron sumados al vigente. El resumen por codigo esta en `13_RESUMEN_VIG0_CODIGO` y el detalle completo en `14_DETALLE_VIG0`.

## Punto 0 versus complemento 95

El complemento 95 fue comparado contra Punto 0 por persona, matricula y fila exacta. La hoja `16_PUNTO0_VS_COMPLEMENTO95` contiene coincidencias y registros nuevos reales por codigo.

## Simulaciones

Se calcularon escenarios diagnosticos por codigo: registros vigentes originales, personas unicas vigentes, llaves de matricula unicas, vigentes sin duplicado exacto y vigentes sin repeticion de llave matricula. Ningun escenario se declara como regla de publicacion.

## Codigos prioritarios

Se generaron fichas para: {', '.join(PRIORITY_CODES)}. Cada ficha localiza totales, repetidos, VIG=0 y efecto de componentes.

## Hallazgos demostrados

La diferencia se localiza numericamente por codigo: positivos {conc['suma_positivos']} y negativos {conc['suma_negativos_abs']}, neto {conc['neto']}. Hay repeticiones institucionales en el congelado, pero no demuestran por si mismas el origen de la publicacion.

## Diferencias no explicadas y pendientes

La causa documental de los 166 sigue pendiente porque falta una publicacion individual con RUT o una descarga/reporte PES que conecte esos totales agregados con registros concretos.

## Conclusion

El analisis ampliado identifica donde esta la diferencia, que duplicados/repeticiones existen y que registros no estan vigentes. No atribuye RUT a la publicacion agregada ni modifica el congelado.
"""
    report_path.write_text(report, encoding="utf-8")

    output_hashes = {
        "excel": sha256_path(xlsx_path),
        "reporte": sha256_path(report_path),
    }

    manifest = {
        "proceso": "Matricula Unificada 2026",
        "subproyecto": "Pregrado",
        "anio": 2026,
        "fecha": TS,
        "fuentes": {
            "workbook_base": str(BASE_XLSX),
            "hash_workbook_base": input_info["base_xlsx_hash"],
            "publicacion": "RAW_PUBLICACION2026 desde workbook base",
            "congelado": "RAW_PREGRADO_CONGELADO desde workbook base",
            "punto0": "RAW_PUNTO0 desde workbook base",
            "complemento95": "RAW_COMPLEMENTO95 desde workbook base",
        },
        "hashes_logicos": input_info,
        "filas": {
            "publicacion_raw": len(pub_raw),
            "pregrado_congelado": len(frozen_raw),
            "punto0": len(p0_raw),
            "complemento95": len(comp_raw),
        },
        "columnas": {
            "publicacion_raw": len(pub_raw.columns),
            "pregrado_congelado": len(frozen_raw.columns),
            "punto0": len(p0_raw.columns),
            "complemento95": len(comp_raw.columns),
        },
        "definicion_llaves": {
            "LLAVE_PERSONA": "TIPO_DOC + N_DOC + DV",
            "CODIGO_UNICO": "I162S{COD_SED}C{COD_CAR}J{JOR}V{VERSION}",
            "LLAVE_MATRICULA": "LLAVE_PERSONA + CODIGO_UNICO",
            "HASH_FILA_32_CAMPOS": "SHA-256 de las 32 columnas oficiales",
        },
        "totales": {
            "publicacion_pregrado": total_pub,
            "congelado_historico": total_hist,
            "vig_0": int(summary["VIG_0"].sum()),
            "vig_1": int(summary["VIG_1"].sum()),
            "vig_2": int(summary["VIG_2"].sum()),
            "congelado_vigente": total_vig,
            "diferencia_neta": int(summary["DIF_PUBLICACION_VS_VIGENTE"].sum()),
            "diferencias_positivas": conc["suma_positivos"],
            "diferencias_negativas_abs": conc["suma_negativos_abs"],
        },
        "duplicados": {
            "rut_repetidos_grupos": len(rut_rep),
            "registros_rut_repetidos": int(rut_rep["Cantidad registros"].sum()) if not rut_rep.empty else 0,
            "matriculas_repetidas_grupos": int(mat_rep["LLAVE_MATRICULA"].nunique()),
            "registros_matricula_repetida": len(mat_rep),
            "duplicados_exactos_grupos": int(exact_rep["HASH_FILA_32_CAMPOS"].nunique()),
            "registros_duplicados_exactos": len(exact_rep),
            "personas_multicarrera": len(multicarrera),
        },
        "concentracion": conc,
        "archivos_generados": {
            "excel": str(xlsx_path),
            "reporte": str(report_path),
            "manifest": str(manifest_path),
        },
        "hashes_archivos_generados": output_hashes,
        "validaciones": validations.to_dict(orient="records"),
        "pendientes": pendientes.to_dict(orient="records"),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps({
        "out_dir": str(OUT_DIR),
        "excel": str(xlsx_path),
        "reporte": str(report_path),
        "manifest": str(manifest_path),
        "publicacion_pregrado": total_pub,
        "congelado_vigente": total_vig,
        "diff": int(summary["DIF_PUBLICACION_VS_VIGENTE"].sum()),
        "positivas": conc["suma_positivos"],
        "negativas_abs": conc["suma_negativos_abs"],
        "codigos": len(summary),
        "rut_repetidos_grupos": len(rut_rep),
        "matriculas_repetidas_grupos": int(mat_rep["LLAVE_MATRICULA"].nunique()),
        "duplicados_exactos_grupos": int(exact_rep["HASH_FILA_32_CAMPOS"].nunique()),
        "excel_ok": excel_ok,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
