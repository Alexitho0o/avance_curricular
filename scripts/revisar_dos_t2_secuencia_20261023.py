"""Verifica secuencia fechada de dos T2 pendientes; no adjudica códigos."""
from pathlib import Path
import json,sys,datetime,collections,subprocess
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R))
from common.diccionario_sies_seguimiento import sha256,DiccionarioSiesSeguimiento
read=lambda p:json.loads(Path(p).read_text());now=datetime.datetime.now(ZoneInfo('America/Santiago'))
O=R/'outputs'/('revision_dos_t2_secuencia_'+now.strftime('%Y%m%d_%H%M%S'));O.mkdir()
def dump(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2,default=str)+'\n')
context=Path('/Users/alexi/.local/share/ipss/cierre_14_jornadas_15_20261023_133759/JORNADAS_15_CONTEXTO_COMPLETO.json')
assert sha256(context)=='cb83605fc892b6fa3879b88b0c4657f58f1b0fd238fa4aca57a46c10e8832892'
rows=read(context);cases={'INF35-03','INF35-04'}
selected=[r for r in rows if r['caso'] in cases and Path(r['fuente']).name=='CUMPLIMIENTO INGRESO CALIFICACIONES.xlsx' and str(r.get('ANO'))=='2024' and str(r.get('PERIODO'))=='2']
assert len(selected)==42
sources={r['fuente']:r['sha256'] for r in selected};sources[str(context)]=sha256(context)
# Relectura de celdas originales: no tomar el derivado como sustituto del archivo.
for p in sorted({r['fuente'] for r in selected}):
 assert sha256(p)==sources[p];w=load_workbook(p,read_only=True,data_only=True)
 for sheet in {r['hoja'] for r in selected if r['fuente']==p}:
  ws=w[sheet];head=next(ws.iter_rows(min_row=2,max_row=2,values_only=True));assert head[3]=='CodCLi' and head[5]=='CODIGO_PLAN_ESTUDIO' and head[28]=='FECHA REPORTE'
  ix={r['fila']:r for r in selected if r['fuente']==p and r['hoja']==sheet};seen=0
  for rn,raw in enumerate(ws.iter_rows(min_row=min(ix),max_row=max(ix),values_only=True),min(ix)):
   if rn not in ix:continue
   r=ix[rn]
   for col,key in [(1,'ANO'),(2,'PERIODO'),(3,'CODCLI'),(5,'CODIGO_PLAN_ESTUDIO'),(9,'CODCARRERA_ALUMNO'),(12,'CODIGO_ASIGNATURA'),(15,'SECCION'),(21,'FECHAINICIO'),(22,'FECHAFIN'),(28,'FECHA REPORTE')]:assert str(raw[col])==str(r[key]),(rn,key)
   seen+=1
  assert seen==len(ix)
 w.close()
refs=[];summary=[]
for case in sorted(cases):
 ss=[r for r in selected if r['caso']==case];dates=sorted({r['FECHA REPORTE'] for r in ss});assert len(dates)==6
 signatures=[]
 for date in dates:
  group=[r for r in ss if r['FECHA REPORTE']==date];sig=sorted({(r['CODIGO_ASIGNATURA'],r['SECCION'],r['CODIGO_PLAN_ESTUDIO']) for r in group});signatures.append(sig)
  summary.append({'caso':case,'fecha_declarada_reporte':date,'asignaturas_secciones':len(sig),'plan':sorted({s[2] for s in sig}),'combinaciones':[{'asignatura':s[0],'seccion':s[1]} for s in sig],'filas':[r['fila'] for r in group],'clasificacion':'B dato observado; C agrupación sin duplicar asignaturas'})
 assert all(s==signatures[0] for s in signatures)
 for r in ss:refs.append({k:r[k] for k in ['caso','fuente','hoja','fila','sha256','FECHA REPORTE','CODIGO_PLAN_ESTUDIO','CODIGO_ASIGNATURA','SECCION','FECHAINICIO','FECHAFIN']})
assert len({(r['caso'],r['CODIGO_ASIGNATURA'],r['SECCION']) for r in selected})==7
# Variación de JORNADA para el mismo contexto histórico, no fechas efectivas.
jrows=[r for r in rows if r['caso'] in cases and str(r.get('ANO'))=='2024' and str(r.get('PERIODO'))=='2' and Path(r['fuente']).name in ['PROMEDIOSDEALUMNOS_2071.xlsx','PROMEDIOSDEALUMNOS_3233.xlsx','PROMEDIOSDEALUMNOS_7804.xlsx','PROMEDIOSDEALUMNOS_12166.xlsx']]
jsummary=[]
for case in sorted(cases):
 for p in sorted({r['fuente'] for r in jrows if r['caso']==case}):
  ss=[r for r in jrows if r['caso']==case and r['fuente']==p];assert len({r['CODCLI'] for r in ss})==1 and {r['PLAN_DE_ESTUDIO'] for r in ss}=={'IINF20241'}
  h=ss[0]['sha256'];assert sha256(p)==h;sources[p]=h
  jsummary.append({'caso':case,'fuente':p,'sha256':h,'hoja':'Hoja1','filas':[r['fila'] for r in ss],'jornada_observada':sorted({r['JORNADA'] for r in ss}),'filas_asignatura':len(ss),'clasificacion':'B; retrospectivo, no fecha efectiva de jornada'})
sm=read(R/'outputs/secciones_cinco_jornadas_20261023_144650/MANIFEST.json');ep=Path(sm['privado'])/'EVIDENCIA_REVISION.json';assert sha256(ep)==sm['privados_hashes'][ep.name];sources[str(ep)]=sha256(ep);ev=read(ep)
targets=[t for t in ev['objetivos'] if t['caso'] in cases and int(t['anio'])==2024 and int(t['periodo'])==2];assert len(targets)==2
A=R/'data/diccionario_sies_seguimiento/activo.json';ah=sha256(A);d=DiccionarioSiesSeguimiento();assert d.edicion=='v7_20261023_inf35_03_t3';dh=d.descriptor['base_privada']['sha256'];result=[]
for t in targets:
 r=json.loads(d.db.execute('select payload from observaciones where id=?',(t['observacion_id'],)).fetchone()[0]);assert not r['codigo_interno_adoptado'] and not r['codigo_documental_atribuible']
 follow=json.loads(d.db.execute('select payload from observaciones where persona=? and codcli=? and plan=? and anio=? and periodo=?',(r['persona'],r['CODCLI'],r['plan'],'2024','3')).fetchone()[0])
 result.append({'caso':t['caso'],'anio':2024,'periodo':2,'codigo_T1':t['codigo_T1'],'codigo_interno_T3':follow['codigo_interno_adoptado'],'T2':'PENDIENTE','atribucion_documental':'PENDIENTE','clasificacion':'B/C/E','nuevas_asignaciones':0})
d.close();assert sha256(A)==ah
for p,h in sources.items():assert sha256(p)==h
dump(O/'SECUENCIA_REPORTES.json',summary);dump(O/'REFERENCIAS.json',refs);dump(O/'JORNADAS_RETROSPECTIVAS.json',jsummary);dump(O/'RESULTADOS.json',result)
report='''# Dos pendientes de jornada T2-2024 tras el cierre INF35-03 T3

Proceso: seguimiento intertrimestral interno IPSS 162. C: revisión focal de dos observaciones, no cálculo de retención ni cambio de reglas generales.

**D/C:** INF35-03 T3-2024 ya está aprobado y aplicado como I162S2C1J1V1 en v7. Su atribución documental permanece E. Esta revisión no extiende esa decisión a T2.

## Nueva comprobación temporal — B/C

Se releen las 42 filas pertinentes de CUMPLIMIENTO INGRESO CALIFICACIONES.xlsx, hoja MASTER INGRESO 2024. El campo FECHA REPORTE distingue seis fechas declaradas: 05-ago, 16-sep, 23-sep, 30-sep, 02-oct y 03-oct de 2024. Son fechas del contenido; las propiedades del libro no certifican cuándo se guardó cada foto. No son seis fuentes independientes.

| Caso | Filas repetidas en seis reportes | Asignaturas/secciones distintas | Evolución observada | Clase |
|---|---:|---:|---|---|
| INF35-03 T2-2024 | 24 | 4 | AT101PCOM/20, AT102PCOM/2, IF103IINF/1 e IF104IINF/1 se mantienen en las seis fechas | B/C |
| INF35-04 T2-2024 | 18 | 3 | AT100PCOM/51, AT101PCOM/1 e IF104IINF/1 se mantienen en las seis fechas | B/C |

B: todas conservan el mismo CODCLI individual y plan IINF20241; sus cursos declaran inicio 01-jul y fin 28-sep-2024. El libro no incluye JORNADA ni VERSION en sus encabezados. C: no se encontró en esta secuencia una sustitución de sección o plan que permita ordenar el cambio. E: estabilidad de las secciones no demuestra jornada administrativa estable; tampoco demuestra que no hubo un cambio.

## Contraste de campos retrospectivos — B/C/E

Para INF35-03, PROMEDIOS2071/3233/7804 muestran D para actividad T2-2024; PROMEDIOS12166 muestra O. Para INF35-04, 2071 muestra V, 3233/7804 D y 12166 O, bajo el mismo plan. Los números de archivo no se usan como fechas efectivas; estas diferencias impiden interpretar JORNADA retrospectiva como historia administrativa inmutable. Los conteos difieren entre extractos y no se suman como estudiantes.

| Caso | Anclas disponibles | Resultado y pendiente preciso | Clase |
|---|---|---|---|
| INF35-03 T2-2024 | T1 documental J2V4; T3 interno aprobado J1V1 | La transición queda entre contextos, pero no ubicada dentro o antes de T2. Las secciones mixtas coexisten, no forman una secuencia. Falta jornada/versión personal atribuible a T2; no heredar V1 desde T3 automáticamente. | B/D/E |
| INF35-04 T2-2024 | T1 documental J1V2; T3 interno aprobado J2V4 | J1V2 sigue candidato de continuidad, no cierre. Contactos de octubre/noviembre describen propuesta diurna e información de cambio posteriores al T2; no fijan jornada efectiva de julio-septiembre. | B/D/E |

E: no hay nuevas correspondencias concluyentes ni se solicita aplicar por mayoría. El cierre T3 de INF35-03 tenía actividad de jornada corroborada en ese contexto; no generaliza a estos T2 con evidencia contextual mixta. El día exacto no es obligatorio para una correspondencia interna, pero sí debe poder atribuirse el contexto al período sin resolver alternativas por supuesto silencioso.

## Control

C: cero cambios durante la revisión. v7 queda con 3.049 observaciones internas pendientes (1.746 personas), 3.213 documentales (1.838 personas). Foco original INF35: 27 reutilizables y ocho pendientes, desglosados en estas dos jornadas y seis homologaciones fuera del bloque. No son tasas de retención. Bases, Excel, índice MU, denominadores y motores intactos. Sin commit/push ni publicación de identificadores.

Referencias exactas, filas y hashes en REFERENCIAS.json y JORNADAS_RETROSPECTIVAS.json; evidencia temporal completa en SECUENCIA_REPORTES.json. Cierre aplicado: ../cierre_inf35_03_t3_20261023_154224/MANIFEST.json.
'''
(O/'INFORME.md').write_text(report);(O/'CHECKPOINT.md').write_text('D/C: v7 conserva INF35-03 T3-2024 J1V1 interno. B/C: secuencia T2 verificada en seis fechas, 42 filas / siete cursos-sección. E: dos T2 siguen pendientes; no hay cambio ordenable de sección/plan ni jornada administrativa atribuida. Se requieren evidencias de asignación por período o una decisión explícita sobre un escenario condicionado; no extrapolar aprobaciones anteriores.\n')
secrets={t['persona'] for t in targets}|{t['CODCLI'] for t in targets}|{t['persona'].split('-')[0] for t in targets}
for p in O.iterdir():assert not any(s in p.read_text() for s in secrets)
qa={'resultado':'OK_REVISION_CON_PENDIENTES_EXPLICITOS','filas_releidas':42,'fechas_reporte':6,'asignaturas_secciones_unicas':7,'personas':2,'nuevas_adjudicaciones':0,'puntero_sin_cambio':sha256(A)==ah,'denominadores_bases_modificados':False,'commit_push':False};dump(O/'QA.json',qa)
dump(O/'MANIFEST.json',{'fecha_hora':now.isoformat(),'clasificacion':'B datos; C revisión; E dos pendientes','fuentes_hashes':sources,'edicion_activa':read(A),'base_privada_hash':dh,'salidas_hashes':{p.name:sha256(p) for p in O.iterdir()},'codigo_hashes':{str(Path(__file__).relative_to(R)):sha256(__file__),'docs/ESPECIFICACION_CIERRE_INF35_03_T3_REVISION_DOS_20261023.md':sha256(R/'docs/ESPECIFICACION_CIERRE_INF35_03_T3_REVISION_DOS_20261023.md')},'QA':qa,'HEAD':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip()})
print(json.dumps({'output':str(O),'QA':qa,'manifest_sha256':sha256(O/'MANIFEST.json')},ensure_ascii=False))
