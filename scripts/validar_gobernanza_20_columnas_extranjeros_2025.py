#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path("/Users/alexi/Documents/GitHub/avance_curricular")
RUN_DIR = ROOT / "estudiantes_extranjeros_2026/resultados/ejecuciones/RECONSTRUCCION_FINAL_EXTRANJEROS_2025_20260625_235455"
PES_FINAL = RUN_DIR / "09_PES/EXTRANJEROS_REGULARES_2025_FINAL_PES_HOJA2.csv"
PRECARGA = ROOT / "estudiantes_extranjeros_2026/Reporte Precarga del Proceso Extranjeros Regulares 2026.csv"
MATRIZ_VIGENCIA = RUN_DIR / "05_VIGENCIA/DECISION_VIGENCIA_FINAL_185_HOJA2.csv"
AUDITORIA_DECISION = RUN_DIR / "05_VIGENCIA/AUDITORIA_DECISION_INSTITUCIONAL_FASE_06.csv"
BASE_EXTRANJEROS = Path("/Users/alexi/Desktop/BASE EXTRANJEROS.xlsx")
PROMEDIOS = ROOT / "PROMEDIOSDEALUMNOS_7804.xlsx"
EXCEL_AUDITORIA_PREVIO = RUN_DIR / "11_REPORTES/12_EXCEL_AUDITORIA_EXTRANJEROS_2025_HOJA2.xlsx"
SCRIPT_GENERADOR = ROOT / "scripts/corregir_fase06_vigencia_extranjeros_2025.py"
SCRIPT_VALIDACION = ROOT / "scripts/validar_gobernanza_20_columnas_extranjeros_2025.py"
CAT_NACIONALIDAD = ROOT / "estudiantes_extranjeros_2026/data/governed/catalogos/NACIONALIDAD_SIES.tsv"
CAT_PAISES = ROOT / "estudiantes_extranjeros_2026/resultados/auditorias/CATALOGO_PAISES_SIES_2026.csv"
GOBERNANZA_NAC = ROOT / "gobernanza_nac.tsv"
CAT_RESIDENCIA = ROOT / "estudiantes_extranjeros_2026/data/governed/catalogos/TIPO_RESIDENCIA_ESTUDIANTE.tsv"
EXPECTED_PES_HASH = "235edaf7363466d8cec6420c00ac8eb56067f3ddd3e82736e33670d2131670b2"

FUNCTIONAL_COLUMNS = [
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "TIPO_RESIDENCIA_ESTUDIANTE",
    "PAIS_DE_ORIGEN",
    "PAIS_ESTUDIOS_SECUNDARIOS",
    "CODIGO_UNICO",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "NOMBRE_UNIVERSIDAD_ORIGEN",
    "PAIS_UNIVERSIDAD_ORIGEN",
    "VIGENCIA",
]

PHYSICAL_BY_FUNCTIONAL = {
    "PAIS_DE_ORIGEN": "PAIS_ORIGEN",
}

STRICT_REQUIRED = {
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "PRIMER_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "PAIS_ESTUDIOS_SECUNDARIOS",
    "CODIGO_UNICO",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "VIGENCIA",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value)
    if text.lower() in {"nan", "none", "<na>", "nat"}:
        return ""
    return text


def trimmed(value: Any) -> str:
    return clean(value).strip()


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_csv(df: pd.DataFrame, path: Path, sep: str = ",") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, sep=sep, quoting=csv.QUOTE_MINIMAL)


def write_json(obj: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            data = df.copy()
            if data.empty:
                data = pd.DataFrame({"SIN_DATOS": [""]})
            data.to_excel(writer, index=False, sheet_name=sheet_name[:31])
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
        for col in ws.columns:
            width = 10
            for cell in list(col)[:150]:
                width = max(width, len(clean(cell.value)) + 2)
            ws.column_dimensions[col[0].column_letter].width = min(width, 70)
    wb.save(path)


def markdown_table(df: pd.DataFrame) -> str:
    columns = [str(col) for col in df.columns]
    rows = [
        "| " + " | ".join(columns) + " |",
        "| " + " | ".join(["---"] * len(columns)) + " |",
    ]
    for _, row in df.fillna("").iterrows():
        values = [clean(row[col]).replace("|", "\\|").replace("\n", " ") for col in df.columns]
        rows.append("| " + " | ".join(values) + " |")
    return "\n".join(rows)


def physical_col(field: str) -> str:
    return PHYSICAL_BY_FUNCTIONAL.get(field, field)


def normalize_value(field: str, value: Any) -> str:
    text = trimmed(value)
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if field in {"TIPO_DOCUMENTO", "DV", "PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES", "SEXO", "CODIGO_UNICO"}:
        return text.strip()
    if field in {
        "NUM_DOCUMENTO",
        "NACIONALIDAD",
        "TIPO_RESIDENCIA_ESTUDIANTE",
        "PAIS_DE_ORIGEN",
        "PAIS_ESTUDIOS_SECUNDARIOS",
        "ANIO_INGRESO_CARRERA_ACTUAL",
        "SEM_INGRESO_CARRERA_ACTUAL",
        "ANIO_INGRESO_CARRERA_ORIGEN",
        "SEM_INGRESO_CARRERA_ORIGEN",
        "PAIS_UNIVERSIDAD_ORIGEN",
        "VIGENCIA",
    }:
        return text.strip()
    return text.strip()


def classify_difference(field: str, prec_value: str, final_value: str, id_registro: str, vig_audit: dict[str, dict[str, str]], matriz_row: dict[str, str]) -> tuple[str, str, str, str, str]:
    norm_prec = normalize_value(field, prec_value)
    norm_final = normalize_value(field, final_value)
    if norm_prec == norm_final:
        if field == "VIGENCIA" and id_registro in vig_audit and vig_audit[id_registro]["TIPO_DECISION"] == "MANTENER_1_DECISION_INSTITUCIONAL":
            return (
                "SIN_DIFERENCIA",
                "DECISION_INSTITUCIONAL_DOCUMENTADA",
                vig_audit[id_registro]["FUENTE_PRINCIPAL"],
                vig_audit[id_registro]["REGLA_APLICADA"],
                "Mantenido en 1 por decision institucional documentada.",
            )
        return ("SIN_DIFERENCIA", "PRECARGA_OFICIAL", "Precarga oficial", "CONSERVAR_VALOR_PRECARGADO", "")
    if field == "FECHA_NACIMIENTO" and normalize_date(norm_prec) == normalize_date(norm_final):
        return ("TRANSFORMACION_TECNICA_NO_MATERIAL", "TRANSFORMACION_TECNICA", "Normalizacion de formato", "TRANSFORMACION_TECNICA_NO_MATERIAL", "")
    if field == "VIGENCIA":
        if id_registro in vig_audit:
            return (
                "CAMBIO_FUNCIONAL_GOBERNADO",
                "DECISION_INSTITUCIONAL_DOCUMENTADA",
                vig_audit[id_registro]["FUENTE_PRINCIPAL"],
                vig_audit[id_registro]["REGLA_APLICADA"],
                vig_audit[id_registro]["OBSERVACION"],
            )
        if trimmed(matriz_row.get("VIGENCIA_FINAL")) == norm_final and trimmed(matriz_row.get("REGLA_APLICADA")):
            return (
                "CAMBIO_FUNCIONAL_GOBERNADO",
                "FUENTE_INSTITUCIONAL",
                trimmed(matriz_row.get("FUENTE_PRINCIPAL")) or "DECISION_VIGENCIA_FINAL_185_HOJA2.csv",
                trimmed(matriz_row.get("REGLA_APLICADA")),
                trimmed(matriz_row.get("OBSERVACION")),
            )
    return ("CAMBIO_NO_GOBERNADO", "PENDIENTE", "", "", "Diferencia contra precarga sin fuente/regla aceptada.")


def normalize_date(value: str) -> str:
    text = trimmed(value)
    if not text:
        return ""
    text = text.split(" ")[0]
    for pattern, repl in [
        (r"^(\d{4})-(\d{2})-(\d{2})$", r"\3-\2-\1"),
        (r"^(\d{2})-(\d{2})-(\d{4})$", r"\1-\2-\3"),
        (r"^(\d{2})/(\d{2})/(\d{4})$", r"\1-\2-\3"),
    ]:
        match = re.match(pattern, text)
        if match:
            return match.expand(repl)
    return text


def is_valid_date(value: str) -> bool:
    text = normalize_date(value)
    try:
        parsed = datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        return False
    return parsed.year <= 2025


def is_integer_text(value: str) -> bool:
    return bool(re.fullmatch(r"\d+", trimmed(value)))


def load_catalogs() -> tuple[set[str], set[str], set[str]]:
    nac = pd.read_csv(CAT_NACIONALIDAD, sep="\t", dtype=str, keep_default_na=False).fillna("")
    res = pd.read_csv(CAT_RESIDENCIA, sep="\t", dtype=str, keep_default_na=False).fillna("")
    allowed_nac = set(nac["CODIGO"].map(trimmed))
    chile_codes = set(nac[nac["ES_CHILE"].map(trimmed).eq("SI")]["CODIGO"].map(trimmed))
    allowed_res = set(res["VALOR"].map(trimmed))
    return allowed_nac, chile_codes, allowed_res


def validate_field(field: str, value: str, row: dict[str, str], allowed_nac: set[str], chile_codes: set[str], allowed_res: set[str]) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    observations: list[str] = []
    v = trimmed(value)
    tipo_doc = trimmed(row.get("TIPO_DOCUMENTO", ""))
    residencia = trimmed(row.get("TIPO_RESIDENCIA_ESTUDIANTE", ""))
    anio_actual = trimmed(row.get("ANIO_INGRESO_CARRERA_ACTUAL", ""))
    anio_origen = trimmed(row.get("ANIO_INGRESO_CARRERA_ORIGEN", ""))
    nombre_univ = trimmed(row.get("NOMBRE_UNIVERSIDAD_ORIGEN", ""))
    pais_univ = trimmed(row.get("PAIS_UNIVERSIDAD_ORIGEN", ""))

    if field in STRICT_REQUIRED and not v:
        errors.append("ERROR_OBLIGATORIEDAD_VACIO")
    if v != clean(value):
        errors.append("ERROR_ESPACIOS_EXTERNOS")
    if field == "TIPO_DOCUMENTO" and v not in {"P", "R"}:
        errors.append("ERROR_DOMINIO_TIPO_DOCUMENTO")
    elif field == "NUM_DOCUMENTO":
        if not re.fullmatch(r"[0-9A-Za-z]+", v):
            errors.append("ERROR_DOCUMENTO_SEPARADORES_NO_AUTORIZADOS")
        if re.search(r"[Ee]\+\d+", v):
            errors.append("ERROR_NOTACION_CIENTIFICA")
    elif field == "DV":
        if tipo_doc == "P" and v:
            observations.append("DV informado en pasaporte; se conserva si proviene de fuente.")
        if tipo_doc == "R" and v and not re.fullmatch(r"[0-9Kk]", v):
            errors.append("ERROR_DOMINIO_DV")
    elif field in {"PRIMER_APELLIDO", "NOMBRES"} and not v:
        errors.append("ERROR_OBLIGATORIEDAD_NOMBRE")
    elif field == "SEXO" and v not in {"M", "H", "NB"}:
        errors.append("ERROR_DOMINIO_SEXO")
    elif field == "FECHA_NACIMIENTO" and (not is_valid_date(v)):
        errors.append("ERROR_FECHA_NACIMIENTO")
    elif field == "NACIONALIDAD":
        if v and (not is_integer_text(v) or v not in allowed_nac):
            errors.append("ERROR_DOMINIO_NACIONALIDAD")
        if v and v in chile_codes:
            errors.append("ERROR_NACIONALIDAD_CHILENA")
    elif field == "TIPO_RESIDENCIA_ESTUDIANTE":
        if v and v not in allowed_res:
            errors.append("ERROR_DOMINIO_RESIDENCIA")
        if not v:
            observations.append("VACIO_AUTORIZADO_POR_REGLA_OFICIAL_ANEXO_V_RESIDENCIA_NO_DISPONIBLE")
    elif field == "PAIS_DE_ORIGEN":
        if residencia in {"2", "3"} and not v:
            errors.append("ERROR_PAIS_ORIGEN_REQUERIDO_POR_RESIDENCIA")
        if residencia == "1" and v:
            errors.append("ERROR_PAIS_ORIGEN_DEBE_VACIARSE_CON_RESIDENCIA_1")
        if not residencia and v:
            errors.append("ERROR_PAIS_ORIGEN_DEBE_VACIARSE_SIN_RESIDENCIA")
        if not residencia and not v:
            observations.append("VACIO_CONDICIONAL_CORRECTO_ANEXO_V_NO_COMPLETAR_PAIS_ORIGEN_SIN_RESIDENCIA")
        if v and (not is_integer_text(v) or not (1 <= int(v) <= 197)):
            errors.append("ERROR_DOMINIO_PAIS_ORIGEN")
        if v and v in chile_codes:
            errors.append("ERROR_PAIS_ORIGEN_CHILE")
    elif field == "PAIS_ESTUDIOS_SECUNDARIOS":
        if v and (not is_integer_text(v) or not (1 <= int(v) <= 197)):
            errors.append("ERROR_DOMINIO_PAIS_ESTUDIOS_SECUNDARIOS")
    elif field == "CODIGO_UNICO" and not v:
        errors.append("ERROR_CODIGO_UNICO_VACIO")
    elif field in {"ANIO_INGRESO_CARRERA_ACTUAL", "ANIO_INGRESO_CARRERA_ORIGEN"}:
        if not re.fullmatch(r"\d{4}", v):
            errors.append("ERROR_FORMATO_ANIO")
        elif int(v) > 2025:
            errors.append("ERROR_ANIO_POSTERIOR_2025")
    elif field in {"SEM_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ORIGEN"}:
        if v not in {"0", "1", "2"}:
            errors.append("ERROR_DOMINIO_SEMESTRE")
    elif field == "ANIO_INGRESO_CARRERA_ORIGEN":
        if anio_actual and anio_origen and is_integer_text(anio_actual) and is_integer_text(anio_origen) and int(anio_origen) > int(anio_actual):
            errors.append("ERROR_ANIO_ORIGEN_POSTERIOR_A_ACTUAL")
    elif field == "NOMBRE_UNIVERSIDAD_ORIGEN":
        if pais_univ and not v:
            errors.append("ERROR_UNIVERSIDAD_ORIGEN_REQUERIDA")
    elif field == "PAIS_UNIVERSIDAD_ORIGEN":
        if nombre_univ and not v:
            errors.append("ERROR_PAIS_UNIVERSIDAD_ORIGEN_REQUERIDO")
        if v and (not is_integer_text(v) or not (1 <= int(v) <= 197)):
            errors.append("ERROR_DOMINIO_PAIS_UNIVERSIDAD_ORIGEN")
    elif field == "VIGENCIA" and v not in {"0", "1"}:
        errors.append("ERROR_DOMINIO_VIGENCIA")
    return errors, observations


def detect_format(path: Path) -> dict[str, Any]:
    data = path.read_bytes()
    text = data.decode("utf-8-sig", errors="replace")
    first_line = text.splitlines()[0] if text.splitlines() else ""
    return {
        "archivo": str(path),
        "bytes": len(data),
        "codificacion_detectada": "utf-8-sig" if data.startswith(b"\xef\xbb\xbf") else "utf-8/ASCII compatible",
        "bom": "SI" if data.startswith(b"\xef\xbb\xbf") else "NO",
        "salto_linea": "CRLF" if b"\r\n" in data else ("LF" if b"\n" in data else "SIN_SALTOS"),
        "delimitador_probable": ";" if first_line.count(";") >= first_line.count(",") else ",",
        "encabezado_probable": "SI" if "TIPO_DOCUMENTO" in first_line and "VIGENCIA" in first_line else "NO",
        "caracteres_control": len(re.findall(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text)),
        "comillas": text.count('"'),
        "comas_internas": text.count(","),
        "saltos_linea_internos": "NO_DETECTADO",
        "notacion_cientifica": len(re.findall(r"\b\d+(?:\.\d+)?[Ee][+-]?\d+\b", text)),
        "valores_decimal_excel": len(re.findall(r"\b\d+\.0\b", text)),
    }


def file_inventory(path: Path, label: str) -> dict[str, Any]:
    stat = path.stat()
    return {
        "ITEM": label,
        "RUTA": str(path),
        "EXISTE": path.exists(),
        "TAMANO_BYTES": stat.st_size,
        "FECHA_MODIFICACION": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "SHA256": sha256(path),
    }


def normalize_text_key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", trimmed(value).upper())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_document_token(value: Any) -> str:
    return re.sub(r"[^0-9A-ZK]", "", trimmed(value).upper())


def final_document_key(row: dict[str, str]) -> str:
    tipo = trimmed(row.get("TIPO_DOCUMENTO", "")).upper()
    num = normalize_document_token(row.get("NUM_DOCUMENTO", ""))
    dv = normalize_document_token(row.get("DV", ""))
    if tipo == "R" and dv:
        return f"{num}{dv}"
    return num


def source_document_key(value: Any) -> str:
    return normalize_document_token(value)


def load_anexo_iii() -> pd.DataFrame:
    if CAT_PAISES.exists():
        src = pd.read_csv(CAT_PAISES, dtype=str, keep_default_na=False).fillna("")
        out = pd.DataFrame({
            "CODIGO_ANEXO_III": src["CODIGO_PAIS"].map(trimmed),
            "NOMBRE_PAIS_ANEXO_III": src["NOMBRE_PAIS"].map(trimmed),
            "NOMBRE_PAIS_NORMALIZADO": src["NOMBRE_PAIS"].map(normalize_text_key),
            "PERMITIDO_NACIONALIDAD": src["PERMITIDO_NACIONALIDAD"].map(trimmed),
            "PERMITIDO_PAIS_ORIGEN": src["PERMITIDO_PAIS_ORIGEN"].map(trimmed),
            "FUENTE": src["FUENTE"].map(trimmed),
        })
    else:
        src = pd.read_csv(CAT_NACIONALIDAD, sep="\t", dtype=str, keep_default_na=False).fillna("")
        out = pd.DataFrame({
            "CODIGO_ANEXO_III": src["CODIGO"].map(trimmed),
            "NOMBRE_PAIS_ANEXO_III": src["DESCRIPCION_OFICIAL"].map(trimmed),
            "NOMBRE_PAIS_NORMALIZADO": src["DESCRIPCION_OFICIAL"].map(normalize_text_key),
            "PERMITIDO_NACIONALIDAD": "SI",
            "PERMITIDO_PAIS_ORIGEN": "SI",
            "FUENTE": src["FUENTE"].map(trimmed),
        })
    return out.sort_values("CODIGO_ANEXO_III", key=lambda s: s.astype(int)).reset_index(drop=True)


def load_gentilicio_mapping(anexo: pd.DataFrame) -> dict[str, dict[str, str]]:
    by_country = {
        normalize_text_key(row["NOMBRE_PAIS_ANEXO_III"]): {
            "CODIGO_ANEXO_III": trimmed(row["CODIGO_ANEXO_III"]),
            "NOMBRE_PAIS_ANEXO_III": trimmed(row["NOMBRE_PAIS_ANEXO_III"]),
            "FUENTE_MAPEO": "01_MAPEO_NACIONALIDADES_ANEXO_III.csv",
        }
        for _, row in anexo.iterrows()
    }
    mapping = dict(by_country)
    if GOBERNANZA_NAC.exists():
        gob = pd.read_csv(GOBERNANZA_NAC, sep="\t", dtype=str, keep_default_na=False).fillna("")
        for _, row in gob.iterrows():
            nac_norm = normalize_text_key(row.get("NACIONALIDAD_NORM", ""))
            pais_norm = normalize_text_key(row.get("PAIS_NORM", ""))
            if not nac_norm or not pais_norm or pais_norm not in by_country:
                continue
            mapping[nac_norm] = {
                **by_country[pais_norm],
                "FUENTE_MAPEO": "gobernanza_nac.tsv + 01_MAPEO_NACIONALIDADES_ANEXO_III.csv",
            }
    return mapping


def institutional_nationality_sources() -> list[dict[str, Any]]:
    sources = []
    base = pd.read_excel(BASE_EXTRANJEROS, sheet_name="Hoja2", dtype=str, keep_default_na=False).fillna("")
    if "NACIONALIDAD" in base.columns and "RUT" in base.columns:
        sources.append({
            "archivo": str(BASE_EXTRANJEROS),
            "archivo_nombre": "BASE EXTRANJEROS.xlsx",
            "hoja": "Hoja2",
            "data": base,
        })
    if PROMEDIOS.exists():
        prom = pd.read_excel(PROMEDIOS, sheet_name="DatosAlumnos", dtype=str, keep_default_na=False).fillna("")
        if "NACIONALIDAD" in prom.columns and "RUT" in prom.columns:
            sources.append({
                "archivo": str(PROMEDIOS),
                "archivo_nombre": "PROMEDIOSDEALUMNOS_7804.xlsx",
                "hoja": "DatosAlumnos",
                "data": prom,
            })
    return sources


def build_source_index(sources: list[dict[str, Any]]) -> dict[str, list[dict[str, str]]]:
    indexed: dict[str, list[dict[str, str]]] = {}
    for source in sources:
        data = source["data"]
        for idx, row in data.iterrows():
            key = source_document_key(row.get("RUT", ""))
            nac = trimmed(row.get("NACIONALIDAD", ""))
            if not key or not nac:
                continue
            indexed.setdefault(key, []).append({
                "NACIONALIDAD_TEXTO_ORIGINAL": nac,
                "NACIONALIDAD_TEXTO_NORMALIZADO": normalize_text_key(nac),
                "ARCHIVO_ORIGEN": source["archivo_nombre"],
                "HOJA_ORIGEN": source["hoja"],
                "FILA_ORIGEN": str(idx + 2),
            })
    return indexed


def resolve_nationality_gaps(final_functional: pd.DataFrame, mapping: dict[str, dict[str, str]]) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, dict[str, str]]]:
    sources = institutional_nationality_sources()
    source_index = build_source_index(sources)
    rows = []
    applicable: dict[str, dict[str, str]] = {}
    blanks = final_functional[final_functional["NACIONALIDAD"].map(trimmed).eq("")]
    for _, rec in blanks.iterrows():
        rec_dict = {col: clean(rec[col]) for col in final_functional.columns}
        id_registro = trimmed(rec["ID_REGISTRO"])
        key = final_document_key(rec_dict)
        candidates = source_index.get(key, [])
        normalized_values = sorted({c["NACIONALIDAD_TEXTO_NORMALIZADO"] for c in candidates if c["NACIONALIDAD_TEXTO_NORMALIZADO"]})
        original_values = sorted({c["NACIONALIDAD_TEXTO_ORIGINAL"] for c in candidates if c["NACIONALIDAD_TEXTO_ORIGINAL"]})
        archivos = sorted({c["ARCHIVO_ORIGEN"] for c in candidates})
        hojas = sorted({c["HOJA_ORIGEN"] for c in candidates})
        filas = sorted({f'{c["ARCHIVO_ORIGEN"]}:{c["HOJA_ORIGEN"]}:{c["FILA_ORIGEN"]}' for c in candidates})
        conflicto = "SI" if len(normalized_values) > 1 else "NO"
        aplicar = "NO"
        codigo = ""
        pais = ""
        observacion = ""
        if not normalized_values:
            observacion = "SIN_NACIONALIDAD_EXPLICITA_EN_FUENTES_AUTORIZADAS"
        elif conflicto == "SI":
            observacion = "CONFLICTO_ENTRE_FUENTES_DE_NACIONALIDAD"
        else:
            norm = normalized_values[0]
            mapped = mapping.get(norm)
            if not mapped:
                observacion = "NACIONALIDAD_SIN_CORRESPONDENCIA_UNICA_EN_ANEXO_III"
            else:
                codigo = mapped["CODIGO_ANEXO_III"]
                pais = mapped["NOMBRE_PAIS_ANEXO_III"]
                if codigo == "38":
                    observacion = "CODIGO_38_CHILE_NO_APLICABLE_A_EXTRANJEROS; NO_SE_APLICA"
                else:
                    aplicar = "SI"
                    observacion = f"CORRESPONDENCIA_GOBERNADA; {mapped['FUENTE_MAPEO']}"
        row = {
            "ID_REGISTRO": id_registro,
            "NUM_DOCUMENTO": trimmed(rec["NUM_DOCUMENTO"]),
            "CODIGO_UNICO": trimmed(rec["CODIGO_UNICO"]),
            "NACIONALIDAD_TEXTO_ORIGINAL": " | ".join(original_values),
            "NACIONALIDAD_TEXTO_NORMALIZADO": " | ".join(normalized_values),
            "CODIGO_ANEXO_III": codigo,
            "NOMBRE_PAIS_ANEXO_III": pais,
            "ARCHIVO_ORIGEN": " | ".join(archivos),
            "HOJA_ORIGEN": " | ".join(hojas),
            "FILA_ORIGEN": " | ".join(filas),
            "CONFLICTO": conflicto,
            "APLICAR": aplicar,
            "OBSERVACION": observacion,
        }
        rows.append(row)
        if aplicar == "SI":
            applicable[id_registro] = row
    recovered = pd.DataFrame(rows)
    problems = recovered[
        recovered["CONFLICTO"].eq("SI")
        | recovered["APLICAR"].ne("SI")
        | recovered["CODIGO_ANEXO_III"].eq("")
        | recovered["CODIGO_ANEXO_III"].eq("38")
    ].copy()
    return recovered, problems, applicable


def classify_difference_corregida(
    field: str,
    prec_value: str,
    final_value: str,
    id_registro: str,
    vig_audit: dict[str, dict[str, str]],
    matriz_row: dict[str, str],
    nac_applicable: dict[str, dict[str, str]],
    row_final: dict[str, str],
) -> tuple[str, str, str, str, str]:
    norm_prec = normalize_value(field, prec_value)
    norm_final = normalize_value(field, final_value)
    if field == "TIPO_RESIDENCIA_ESTUDIANTE" and not norm_final:
        return (
            "VACIO_AUTORIZADO_POR_REGLA_OFICIAL",
            "REGLA_OFICIAL_SIN_INFORMACION",
            "REGLAS_VALIDACION_INSTRUCTIVO 2.md",
            "ANEXO_V_RESIDENCIA_NO_DISPONIBLE",
            "No se completa TIPO_RESIDENCIA_ESTUDIANTE por falta de informacion institucional suficiente.",
        )
    if field == "PAIS_DE_ORIGEN" and not norm_final and not trimmed(row_final.get("TIPO_RESIDENCIA_ESTUDIANTE", "")):
        return (
            "VACIO_CONDICIONAL_CORRECTO",
            "REGLA_OFICIAL_CONDICIONAL",
            "REGLAS_VALIDACION_INSTRUCTIVO 2.md",
            "ANEXO_V_NO_COMPLETAR_PAIS_ORIGEN_SIN_RESIDENCIA",
            "No se completa PAIS_DE_ORIGEN porque TIPO_RESIDENCIA_ESTUDIANTE no esta disponible.",
        )
    if field == "NACIONALIDAD" and norm_prec != norm_final and id_registro in nac_applicable:
        audit = nac_applicable[id_registro]
        return (
            "CAMBIO_FUNCIONAL_GOBERNADO",
            "FUENTE_INSTITUCIONAL",
            audit["ARCHIVO_ORIGEN"],
            "NACIONALIDAD_DESDE_FUENTE_INSTITUCIONAL_Y_ANEXO_III",
            f"{audit['NACIONALIDAD_TEXTO_ORIGINAL']} -> {audit['CODIGO_ANEXO_III']} {audit['NOMBRE_PAIS_ANEXO_III']}",
        )
    return classify_difference(field, prec_value, final_value, id_registro, vig_audit, matriz_row)


def governance_for(
    final_functional: pd.DataFrame,
    precarga: pd.DataFrame,
    prec_functional: pd.DataFrame,
    matriz_by_id: dict[str, dict[str, str]],
    vig_audit: dict[str, dict[str, str]],
    nac_applicable: dict[str, dict[str, str]],
    allowed_nac: set[str],
    chile_codes: set[str],
    allowed_res: set[str],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    governance_rows: list[dict[str, Any]] = []
    comparison_rows: list[dict[str, Any]] = []
    difference_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []

    for idx in range(len(final_functional)):
        id_registro = final_functional.loc[idx, "ID_REGISTRO"]
        fila = str(final_functional.loc[idx, "FILA_PRECARGA"])
        matriz_row = matriz_by_id.get(id_registro, {})
        row_final = {field: final_functional.loc[idx, field] for field in FUNCTIONAL_COLUMNS}
        row_comp: dict[str, Any] = {
            "ID_REGISTRO": id_registro,
            "FILA_PRECARGA": fila,
            "CODIGO_IES_NUM": precarga.loc[idx, "CODIGO_IES_NUM"],
            "TIPO_DOCUMENTO": final_functional.loc[idx, "TIPO_DOCUMENTO"],
            "NUM_DOCUMENTO": final_functional.loc[idx, "NUM_DOCUMENTO"],
            "CODIGO_UNICO": final_functional.loc[idx, "CODIGO_UNICO"],
        }
        for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1):
            prec_val = clean(prec_functional.loc[idx, field])
            final_val = clean(final_functional.loc[idx, field])
            norm_prec = normalize_value(field, prec_val)
            norm_final = normalize_value(field, final_val)
            tipo_cambio, origen, fuente, regla, obs_diff = classify_difference_corregida(
                field, prec_val, final_val, id_registro, vig_audit, matriz_row, nac_applicable, row_final
            )
            errors, obs_control = validate_field(field, final_val, row_final, allowed_nac, chile_codes, allowed_res)
            coincide = "SI" if norm_prec == norm_final else "NO"
            modificado = "NO" if coincide == "SI" else "SI"
            blocking_error = any(
                e.startswith("ERROR_DOMINIO")
                or e.startswith("ERROR_OBLIGATORIEDAD")
                or e.startswith("ERROR_NACIONALIDAD")
                or e.startswith("ERROR_PAIS")
                or e.startswith("ERROR_UNIVERSIDAD")
                or e.startswith("ERROR_ANIO")
                or e.startswith("ERROR_FECHA")
                or e.startswith("ERROR_CODIGO")
                for e in errors
            )
            gobernado = "SI" if tipo_cambio != "CAMBIO_NO_GOBERNADO" and not blocking_error else "NO"
            resultado = "OK" if gobernado == "SI" and not errors else "ERROR"
            if field == "TIPO_RESIDENCIA_ESTUDIANTE" and not norm_final and not errors:
                resultado = "VACIO_AUTORIZADO_POR_REGLA_OFICIAL"
            if field == "PAIS_DE_ORIGEN" and not norm_final and not trimmed(row_final.get("TIPO_RESIDENCIA_ESTUDIANTE", "")) and not errors:
                resultado = "VACIO_CONDICIONAL_CORRECTO"
            row = {
                "ID_REGISTRO": id_registro,
                "FILA_PRECARGA": fila,
                "CODIGO_IES_NUM": precarga.loc[idx, "CODIGO_IES_NUM"],
                "TIPO_DOCUMENTO": final_functional.loc[idx, "TIPO_DOCUMENTO"],
                "NUM_DOCUMENTO": final_functional.loc[idx, "NUM_DOCUMENTO"],
                "CODIGO_UNICO": final_functional.loc[idx, "CODIGO_UNICO"],
                "POSICION_SIES": pos,
                "CAMPO": field,
                "VALOR_PRECARGA": prec_val,
                "VALOR_FINAL": final_val,
                "VALOR_NORMALIZADO_PRECARGA": norm_prec,
                "VALOR_NORMALIZADO_FINAL": norm_final,
                "COINCIDE_PRECARGA": coincide,
                "FUE_MODIFICADO": modificado,
                "ORIGEN_VALOR_FINAL": origen,
                "FUENTE_RESPALDO": fuente,
                "REGLA_APLICADA": regla,
                "NIVEL_RESPALDO": "ALTO" if origen in {"PRECARGA_OFICIAL", "DECISION_INSTITUCIONAL_DOCUMENTADA", "FUENTE_INSTITUCIONAL", "REGLA_OFICIAL_SIN_INFORMACION", "REGLA_OFICIAL_CONDICIONAL"} else ("MEDIO" if origen == "TRANSFORMACION_TECNICA" else "PENDIENTE"),
                "GOBERNADO": gobernado,
                "RESULTADO_CONTROL": resultado,
                "OBSERVACION": " | ".join([obs_diff] + obs_control + errors).strip(" |"),
            }
            governance_rows.append(row)
            row_comp[f"{field}_PRECARGA"] = prec_val
            row_comp[f"{field}_FINAL"] = final_val
            row_comp[f"{field}_IGUAL"] = coincide
            row_comp[f"{field}_TIPO_CAMBIO"] = tipo_cambio
            row_comp[f"{field}_FUENTE"] = fuente
            row_comp[f"{field}_REGLA"] = regla
            row_comp[f"{field}_GOBERNADO"] = gobernado
            if tipo_cambio not in {"SIN_DIFERENCIA"}:
                difference_rows.append({
                    "ID_REGISTRO": id_registro,
                    "FILA_PRECARGA": fila,
                    "CAMPO": field,
                    "VALOR_PRECARGA": prec_val,
                    "VALOR_FINAL": final_val,
                    "TIPO_DIFERENCIA": tipo_cambio,
                    "FUENTE_RESPALDO": fuente,
                    "REGLA_APLICADA": regla,
                    "GOBERNADO": gobernado,
                    "OBSERVACION": obs_diff,
                })
            for err in errors:
                error_rows.append({
                    "ID_REGISTRO": id_registro,
                    "FILA_PRECARGA": fila,
                    "CAMPO": field,
                    "VALOR_FINAL": final_val,
                    "ERROR": err,
                    "BLOQUEANTE": "SI",
                })
        comparison_rows.append(row_comp)

    governance = pd.DataFrame(governance_rows)
    comparison = pd.DataFrame(comparison_rows)
    differences = pd.DataFrame(difference_rows)
    errors = pd.DataFrame(error_rows)
    if errors.empty:
        errors = pd.DataFrame(columns=["ID_REGISTRO", "FILA_PRECARGA", "CAMPO", "VALOR_FINAL", "ERROR", "BLOQUEANTE"])
    if differences.empty:
        differences = pd.DataFrame(columns=["ID_REGISTRO", "FILA_PRECARGA", "CAMPO", "VALOR_PRECARGA", "VALOR_FINAL", "TIPO_DIFERENCIA", "FUENTE_RESPALDO", "REGLA_APLICADA", "GOBERNADO", "OBSERVACION"])

    resumen_col_rows = []
    for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1):
        g = governance[governance["CAMPO"].eq(field)]
        resumen_col_rows.append({
            "POSICION": pos,
            "COLUMNA": field,
            "FILAS": len(g),
            "VACIOS": int(g["VALOR_FINAL"].map(trimmed).eq("").sum()),
            "NO_VACIOS": int(g["VALOR_FINAL"].map(trimmed).ne("").sum()),
            "VALORES_UNICOS": int(g["VALOR_FINAL"].nunique()),
            "COINCIDENCIAS_CON_PRECARGA": int(g["COINCIDE_PRECARGA"].eq("SI").sum()),
            "DIFERENCIAS": int(g["COINCIDE_PRECARGA"].eq("NO").sum()),
            "CAMBIOS_GOBERNADOS": int(((g["COINCIDE_PRECARGA"].eq("NO")) & (g["GOBERNADO"].eq("SI"))).sum()),
            "TRANSFORMACIONES_TECNICAS": int(g["REGLA_APLICADA"].eq("TRANSFORMACION_TECNICA_NO_MATERIAL").sum()),
            "CAMBIOS_NO_GOBERNADOS": int(((g["COINCIDE_PRECARGA"].eq("NO")) & (g["GOBERNADO"].eq("NO"))).sum()),
            "ERRORES_DOMINIO": int(g["OBSERVACION"].str.contains("ERROR_DOMINIO|ERROR_NACIONALIDAD_CHILENA|ERROR_PAIS_ORIGEN_CHILE|ERROR_FECHA|ERROR_ANIO|ERROR_SEMESTRE", regex=True, na=False).sum()),
            "ERRORES_OBLIGATORIEDAD": int(g["OBSERVACION"].str.contains("ERROR_OBLIGATORIEDAD|ERROR_.*REQUERID", regex=True, na=False).sum()),
            "RESULTADO": "OK" if not g["RESULTADO_CONTROL"].eq("ERROR").any() else "ERROR",
        })
    resumen_col = pd.DataFrame(resumen_col_rows)
    return governance, comparison, differences, errors, resumen_col


def write_hashes(out_dir: Path) -> pd.DataFrame:
    hash_rows = []
    for path in sorted(out_dir.glob("*")):
        if path.is_file() and path.name != "HASHES_VALIDACION_FINAL.sha256":
            hash_rows.append(f"{sha256(path)}  {path}")
    (out_dir / "HASHES_VALIDACION_FINAL.sha256").write_text("\n".join(hash_rows) + "\n", encoding="utf-8")
    return pd.DataFrame([
        {"ARCHIVO": str(path), "SHA256": sha256(path)}
        for path in sorted(out_dir.glob("*"))
        if path.is_file()
    ])


def main_corregido() -> int:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = RUN_DIR / "12_VALIDACION_FINAL_SIES" / f"CORRECCION_DATOS_PERSONALES_{timestamp}"
    out_dir.mkdir(parents=True, exist_ok=False)

    pes_hash = sha256(PES_FINAL)
    inventory_items = [
        file_inventory(PES_FINAL, "archivo PES generado revisado"),
        file_inventory(PRECARGA, "precarga oficial"),
        file_inventory(MATRIZ_VIGENCIA, "matriz final de vigencia"),
        file_inventory(AUDITORIA_DECISION, "auditoria decision institucional"),
        file_inventory(BASE_EXTRANJEROS, "BASE EXTRANJEROS.xlsx"),
        file_inventory(PROMEDIOS, "PROMEDIOSDEALUMNOS_7804.xlsx"),
        file_inventory(SCRIPT_VALIDACION, "script validacion/correccion"),
        file_inventory(CAT_NACIONALIDAD, "catalogo nacionalidad SIES"),
        file_inventory(CAT_PAISES, "catalogo paises Anexo III local"),
    ]
    inventory = pd.DataFrame(inventory_items)
    if pes_hash != EXPECTED_PES_HASH:
        summary = {
            "estado_final": "BLOQUEADO_POR_CAMBIO_DEL_ARCHIVO_PES",
            "hash_esperado": EXPECTED_PES_HASH,
            "hash_obtenido": pes_hash,
        }
        write_csv(inventory, out_dir / "INVENTARIO_CONGELAMIENTO.csv")
        write_json(summary, out_dir / "RESUMEN_VALIDACION_FINAL.json")
        return 2

    precarga = pd.read_csv(PRECARGA, sep=";", encoding="cp1252", dtype=str, keep_default_na=False).fillna("")
    final = pd.read_csv(PES_FINAL, sep=";", dtype=str, keep_default_na=False).fillna("")
    matriz = pd.read_csv(MATRIZ_VIGENCIA, dtype=str, keep_default_na=False).fillna("")
    audit = pd.read_csv(AUDITORIA_DECISION, dtype=str, keep_default_na=False).fillna("")
    allowed_nac, chile_codes, allowed_res = load_catalogs()
    vig_audit = {trimmed(r["ID_REGISTRO"]): {k: trimmed(v) for k, v in r.items()} for _, r in audit.iterrows()}
    matriz_by_id = {trimmed(r["ID_REGISTRO"]): {k: trimmed(v) for k, v in r.items()} for _, r in matriz.iterrows()}

    final_functional = pd.DataFrame()
    prec_functional = pd.DataFrame()
    for field in FUNCTIONAL_COLUMNS:
        final_functional[field] = final[physical_col(field)]
        prec_functional[field] = precarga[physical_col(field)]
    final_functional.insert(0, "ID_REGISTRO", [f"PRECARGA_{i:04d}" for i in range(1, len(final_functional) + 1)])
    final_functional.insert(1, "FILA_PRECARGA", list(range(1, len(final_functional) + 1)))
    prec_out = precarga.copy()
    prec_out.insert(0, "ID_REGISTRO", [f"PRECARGA_{i:04d}" for i in range(1, len(prec_out) + 1)])
    prec_out.insert(1, "FILA_PRECARGA", list(range(1, len(prec_out) + 1)))

    anexo = load_anexo_iii()
    mapping = load_gentilicio_mapping(anexo)
    recovered, nac_problems, nac_applicable = resolve_nationality_gaps(final_functional, mapping)
    corrected = final_functional.copy()
    for id_registro, audit_row in nac_applicable.items():
        corrected.loc[corrected["ID_REGISTRO"].eq(id_registro), "NACIONALIDAD"] = audit_row["CODIGO_ANEXO_III"]

    residencia_audit_rows = []
    for _, row in corrected.iterrows():
        residencia_audit_rows.append({
            "ID_REGISTRO": row["ID_REGISTRO"],
            "FILA_PRECARGA": row["FILA_PRECARGA"],
            "NUM_DOCUMENTO": row["NUM_DOCUMENTO"],
            "CODIGO_UNICO": row["CODIGO_UNICO"],
            "TIPO_RESIDENCIA_ESTUDIANTE": row["TIPO_RESIDENCIA_ESTUDIANTE"],
            "PAIS_DE_ORIGEN": row["PAIS_DE_ORIGEN"],
            "ORIGEN_TIPO_RESIDENCIA": "REGLA_OFICIAL_SIN_INFORMACION",
            "REGLA_TIPO_RESIDENCIA": "ANEXO_V_RESIDENCIA_NO_DISPONIBLE",
            "RESULTADO_TIPO_RESIDENCIA": "VACIO_AUTORIZADO_POR_REGLA_OFICIAL" if not trimmed(row["TIPO_RESIDENCIA_ESTUDIANTE"]) else "INFORMADO",
            "ORIGEN_PAIS_DE_ORIGEN": "REGLA_OFICIAL_CONDICIONAL",
            "REGLA_PAIS_DE_ORIGEN": "ANEXO_V_NO_COMPLETAR_PAIS_ORIGEN_SIN_RESIDENCIA",
            "RESULTADO_PAIS_DE_ORIGEN": "VACIO_CONDICIONAL_CORRECTO" if not trimmed(row["TIPO_RESIDENCIA_ESTUDIANTE"]) and not trimmed(row["PAIS_DE_ORIGEN"]) else "REVISAR_COMBINACION",
        })
    residencia_audit = pd.DataFrame(residencia_audit_rows)

    governance, comparison, differences, errors, resumen_col = governance_for(
        corrected,
        precarga,
        prec_functional,
        matriz_by_id,
        vig_audit,
        nac_applicable,
        allowed_nac,
        chile_codes,
        allowed_res,
    )

    duplicate_persona_carrera = int(corrected.duplicated(subset=["TIPO_DOCUMENTO", "NUM_DOCUMENTO", "CODIGO_UNICO"], keep=False).sum())
    solo_precarga = 0 if len(precarga) == len(corrected) else max(len(precarga) - len(corrected), 0)
    solo_final = 0 if len(precarga) == len(corrected) else max(len(corrected) - len(precarga), 0)
    cambios_no_gobernados = int(differences["TIPO_DIFERENCIA"].eq("CAMBIO_NO_GOBERNADO").sum())
    errores_bloqueantes = int(errors["BLOQUEANTE"].eq("SI").sum())
    nacionalidad_vacia = int(corrected["NACIONALIDAD"].map(trimmed).eq("").sum())
    nacionalidad_codigo_38 = int(corrected["NACIONALIDAD"].map(trimmed).eq("38").sum())
    conflictos_nacionalidad = int(recovered["CONFLICTO"].eq("SI").sum()) if not recovered.empty else 0
    nacionalidad_sin_correspondencia = int(recovered["CODIGO_ANEXO_III"].map(trimmed).eq("").sum()) if not recovered.empty else 0
    nacionalidad_no_aplicada_38 = int(recovered["CODIGO_ANEXO_III"].eq("38").sum()) if not recovered.empty else 0
    tipo_residencia_vacios = int(corrected["TIPO_RESIDENCIA_ESTUDIANTE"].map(trimmed).eq("").sum())
    pais_origen_vacios = int(corrected["PAIS_DE_ORIGEN"].map(trimmed).eq("").sum())

    can_make_ready = (
        len(corrected) == 185
        and len(FUNCTIONAL_COLUMNS) == 20
        and len(governance) == 3700
        and cambios_no_gobernados == 0
        and conflictos_nacionalidad == 0
        and nacionalidad_vacia == 0
        and nacionalidad_codigo_38 == 0
        and nacionalidad_sin_correspondencia == 0
        and errores_bloqueantes == 0
        and duplicate_persona_carrera == 0
        and solo_precarga == 0
        and solo_final == 0
        and int(corrected["VIGENCIA"].eq("0").sum()) == 24
        and int(corrected["VIGENCIA"].eq("1").sum()) == 161
        and int((~corrected["VIGENCIA"].isin(["0", "1"])).sum()) == 0
        and int(corrected["CODIGO_UNICO"].map(trimmed).eq("").sum()) == 0
    )

    pes_ready_path = ""
    pes_ready_hash = ""
    desktop_path = ""
    desktop_hash = ""
    desktop_identical = "NO_APLICA"
    ready_format = {
        "archivo": "",
        "filas_relectura": "",
        "columnas_relectura": "",
        "encabezado": "NO_APLICA",
        "delimitador": "NO_APLICA",
        "codificacion": "UTF-8",
        "bom": "NO_APLICA",
    }
    if can_make_ready:
        ready_name = f"EXTRANJEROS_REGULARES_2025_PES_READY_{timestamp}.csv"
        ready = out_dir / ready_name
        with ready.open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=",", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
            for _, row in corrected[FUNCTIONAL_COLUMNS].iterrows():
                writer.writerow([clean(row[col]) for col in FUNCTIONAL_COLUMNS])
        reread = pd.read_csv(ready, header=None, sep=",", dtype=str, keep_default_na=False).fillna("")
        fmt_ready = detect_format(ready)
        ready_format = {
            "archivo": str(ready),
            "filas_relectura": int(reread.shape[0]),
            "columnas_relectura": int(reread.shape[1]),
            "encabezado": fmt_ready["encabezado_probable"],
            "delimitador": fmt_ready["delimitador_probable"],
            "codificacion": fmt_ready["codificacion_detectada"],
            "bom": fmt_ready["bom"],
        }
        if (
            reread.shape == (185, 20)
            and reread.iloc[:, 8].map(trimmed).ne("").all()
            and reread.iloc[:, 9].map(trimmed).eq("").sum() == 185
            and reread.iloc[:, 10].map(trimmed).eq("").sum() == 185
            and reread.iloc[:, 19].eq("0").sum() == 24
            and reread.iloc[:, 19].eq("1").sum() == 161
            and fmt_ready["encabezado_probable"] == "NO"
            and fmt_ready["delimitador_probable"] == ","
        ):
            desktop = Path.home() / "Desktop" / ready_name
            shutil.copy2(ready, desktop)
            pes_ready_path = str(ready)
            desktop_path = str(desktop)
            pes_ready_hash = sha256(ready)
            desktop_hash = sha256(desktop)
            desktop_identical = "SI" if pes_ready_hash == desktop_hash else "NO"
        else:
            ready.unlink(missing_ok=True)
            can_make_ready = False

    if nacionalidad_vacia or nacionalidad_codigo_38 or conflictos_nacionalidad or nacionalidad_sin_correspondencia:
        status = "BLOQUEADO_POR_NACIONALIDAD_INCOMPLETA"
    elif can_make_ready and desktop_identical == "SI":
        status = "LISTO_PARA_CARGA_PES"
    else:
        status = "BLOQUEADO_POR_VALIDACION_FINAL"

    pruebas_fallidas = []
    if len(governance) != 3700:
        pruebas_fallidas.append("CONTROLES_GOBERNANZA_DISTINTO_3700")
    if cambios_no_gobernados:
        pruebas_fallidas.append("CAMBIOS_NO_GOBERNADOS")
    if errores_bloqueantes:
        pruebas_fallidas.append("ERRORES_BLOQUEANTES")
    if conflictos_nacionalidad:
        pruebas_fallidas.append("CONFLICTOS_NACIONALIDAD")
    if nacionalidad_sin_correspondencia:
        pruebas_fallidas.append("NACIONALIDAD_SIN_CORRESPONDENCIA_ANEXO_III")
    if nacionalidad_no_aplicada_38:
        pruebas_fallidas.append("NACIONALIDAD_CODIGO_38_NO_APLICABLE")
    if nacionalidad_vacia:
        pruebas_fallidas.append("NACIONALIDAD_VACIA")
    if duplicate_persona_carrera:
        pruebas_fallidas.append("DUPLICADOS_PERSONA_CARRERA")
    if can_make_ready and desktop_identical != "SI":
        pruebas_fallidas.append("COPIA_ESCRITORIO_NO_IDENTICA")

    formatos = pd.DataFrame([detect_format(PES_FINAL), ready_format])
    dictionary = pd.DataFrame([
        {
            "POSICION": pos,
            "COLUMNA": field,
            "COLUMNA_FISICA_PRECARGA": physical_col(field),
            "COLUMNA_FISICA_FINAL": physical_col(field),
            "OBLIGATORIEDAD_CONTROLADA": "SI" if field in STRICT_REQUIRED else "CONDICIONAL/O_PERMITIDO_SEGUN_FUENTE",
        }
        for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1)
    ])

    resumen = {
        "proceso": "Estudiantes Extranjeros Regulares 2025 - SIES 2026",
        "anio_datos": 2025,
        "estado_final": status,
        "carpeta_validacion": str(out_dir),
        "archivo_final_revisado": str(PES_FINAL),
        "hash_archivo_final_revisado": pes_hash,
        "hash_esperado_archivo_final": EXPECTED_PES_HASH,
        "precarga_utilizada": str(PRECARGA),
        "hash_precarga": sha256(PRECARGA),
        "base_institucional": str(BASE_EXTRANJEROS),
        "hash_base_institucional": sha256(BASE_EXTRANJEROS),
        "fuente_secundaria_nacionalidad": str(PROMEDIOS),
        "hash_fuente_secundaria_nacionalidad": sha256(PROMEDIOS),
        "filas": int(len(corrected)),
        "columnas": 20,
        "controles_gobernanza_esperados": 3700,
        "controles_gobernanza_ejecutados": int(len(governance)),
        "nacionalidades_encontradas_fuente_institucional": int(recovered["NACIONALIDAD_TEXTO_NORMALIZADO"].map(trimmed).ne("").sum()) if not recovered.empty else 0,
        "nacionalidades_mapeadas_al_anexo_iii": int(recovered["CODIGO_ANEXO_III"].map(trimmed).ne("").sum()) if not recovered.empty else 0,
        "nacionalidades_aplicadas": int(recovered["APLICAR"].eq("SI").sum()) if not recovered.empty else 0,
        "nacionalidades_sin_correspondencia": nacionalidad_sin_correspondencia,
        "nacionalidades_no_aplicadas_codigo_38": nacionalidad_no_aplicada_38,
        "conflictos": conflictos_nacionalidad,
        "tipo_residencia_vacios": tipo_residencia_vacios,
        "pais_de_origen_vacios": pais_origen_vacios,
        "regla_oficial_aplicada": "ANEXO_V_RESIDENCIA_NO_DISPONIBLE; ANEXO_V_NO_COMPLETAR_PAIS_ORIGEN_SIN_RESIDENCIA",
        "coincidencias_exactas": int(governance["COINCIDE_PRECARGA"].eq("SI").sum()),
        "transformaciones_tecnicas": int(differences["TIPO_DIFERENCIA"].eq("TRANSFORMACION_TECNICA_NO_MATERIAL").sum()) if not differences.empty else 0,
        "cambios_funcionales_gobernados": int(differences["TIPO_DIFERENCIA"].eq("CAMBIO_FUNCIONAL_GOBERNADO").sum()) if not differences.empty else 0,
        "cambios_no_gobernados": cambios_no_gobernados,
        "errores_bloqueantes": errores_bloqueantes,
        "duplicados_persona_carrera": duplicate_persona_carrera,
        "solo_precarga": solo_precarga,
        "solo_final": solo_final,
        "vigencia_0": int(corrected["VIGENCIA"].eq("0").sum()),
        "vigencia_1": int(corrected["VIGENCIA"].eq("1").sum()),
        "vigencia_fuera_dominio": int((~corrected["VIGENCIA"].isin(["0", "1"])).sum()),
        "codigo_unico_vacio": int(corrected["CODIGO_UNICO"].map(trimmed).eq("").sum()),
        "nacionalidad_vacia": nacionalidad_vacia,
        "encabezado": ready_format["encabezado"],
        "delimitador": ready_format["delimitador"],
        "archivo_pes_ready": pes_ready_path,
        "hash_pes_ready": pes_ready_hash,
        "copia_escritorio": desktop_path,
        "hash_copia_escritorio": desktop_hash,
        "copia_escritorio_identica": desktop_identical,
        "pruebas_fallidas": pruebas_fallidas,
    }

    write_csv(inventory, out_dir / "INVENTARIO_CONGELAMIENTO.csv")
    write_csv(anexo, out_dir / "01_MAPEO_NACIONALIDADES_ANEXO_III.csv")
    write_csv(recovered, out_dir / "02_NACIONALIDAD_104_RECUPERADA.csv")
    write_csv(nac_problems, out_dir / "03_NACIONALIDAD_CONFLICTOS_O_NO_MAPEADA.csv")
    write_csv(corrected[FUNCTIONAL_COLUMNS], out_dir / "14_EXTRANJEROS_REGULARES_2025_COMPLETO_20_COLUMNAS.csv")
    write_csv(recovered, out_dir / "15_AUDITORIA_APLICACION_NACIONALIDAD.csv")
    write_csv(residencia_audit, out_dir / "16_AUDITORIA_REGLA_RESIDENCIA_PAIS_ORIGEN.csv")
    write_csv(governance, out_dir / "MATRIZ_GOBERNANZA_185_X_20.csv")
    write_csv(comparison, out_dir / "COMPARACION_PRECARGA_VS_FINAL_20_COLUMNAS.csv")
    write_csv(differences, out_dir / "DIFERENCIAS_GOBERNANZA_20_COLUMNAS.csv")
    write_csv(errors, out_dir / "ERRORES_VALIDACION_FINAL.csv")
    write_csv(resumen_col, out_dir / "RESUMEN_POR_COLUMNA.csv")
    write_csv(formatos, out_dir / "FORMATOS.csv")
    write_json(resumen, out_dir / "RESUMEN_VALIDACION_FINAL.json")

    report_lines = [
        "# Correccion final de completitud y validacion PES",
        "",
        f"Estado final: {status}",
        f"Archivo final revisado: {PES_FINAL}",
        f"Hash archivo final revisado: {pes_hash}",
        f"Nacionalidades encontradas: {resumen['nacionalidades_encontradas_fuente_institucional']}",
        f"Nacionalidades mapeadas al Anexo III: {resumen['nacionalidades_mapeadas_al_anexo_iii']}",
        f"Nacionalidades aplicadas: {resumen['nacionalidades_aplicadas']}",
        f"Nacionalidades no aplicadas por codigo 38: {nacionalidad_no_aplicada_38}",
        f"Conflictos: {conflictos_nacionalidad}",
        f"TIPO_RESIDENCIA vacios: {tipo_residencia_vacios}",
        f"PAIS_DE_ORIGEN vacios: {pais_origen_vacios}",
        f"Controles ejecutados: {len(governance)}",
        f"Cambios no gobernados: {cambios_no_gobernados}",
        f"Errores bloqueantes: {errores_bloqueantes}",
        f"VIGENCIA 0: {resumen['vigencia_0']}",
        f"VIGENCIA 1: {resumen['vigencia_1']}",
        f"PES READY: {pes_ready_path or 'NO_GENERADO'}",
        "",
        "## Resultado por columna",
        "",
        markdown_table(resumen_col),
    ]
    (out_dir / "REPORTE_VALIDACION_FINAL.md").write_text("\n".join(report_lines) + "\n", encoding="utf-8")

    hashes = write_hashes(out_dir)
    write_excel(out_dir / "AUDITORIA_GOBERNANZA_20_COLUMNAS_Y_FORMATO_SIES.xlsx", {
        "RESUMEN": pd.DataFrame([resumen]),
        "DICCIONARIO_20_COLUMNAS": dictionary,
        "RESUMEN_POR_COLUMNA": resumen_col,
        "GOBERNANZA_3700": governance,
        "PRECARGA_185": prec_out,
        "FINAL_185": corrected,
        "COMPARACION_20_COLUMNAS": comparison,
        "DIFERENCIAS": differences,
        "ERRORES": errors,
        "VIGENCIA": matriz,
        "NACIONALIDAD": recovered,
        "RESIDENCIA_PAIS_ORIGEN": residencia_audit,
        "FORMATOS": formatos,
        "HASHES": hashes,
    })
    write_hashes(out_dir)

    print(json.dumps(resumen, ensure_ascii=False, indent=2))
    return 0 if status == "LISTO_PARA_CARGA_PES" else 2


def main() -> int:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_dir = RUN_DIR / "12_VALIDACION_FINAL_SIES" / f"VALIDACION_GOBERNANZA_20_COLUMNAS_{timestamp}"
    out_dir.mkdir(parents=True, exist_ok=False)

    pes_hash = sha256(PES_FINAL)
    inventory_items = [
        file_inventory(PES_FINAL, "archivo PES generado"),
        file_inventory(PRECARGA, "precarga oficial"),
        file_inventory(MATRIZ_VIGENCIA, "matriz final de vigencia"),
        file_inventory(AUDITORIA_DECISION, "auditoria decision institucional"),
        file_inventory(BASE_EXTRANJEROS, "BASE EXTRANJEROS.xlsx"),
        file_inventory(EXCEL_AUDITORIA_PREVIO, "Excel auditoria previo"),
        file_inventory(SCRIPT_GENERADOR, "script generador"),
        file_inventory(SCRIPT_VALIDACION, "script validacion final"),
    ]
    inventory = pd.DataFrame(inventory_items)
    if pes_hash != EXPECTED_PES_HASH:
        summary = {
            "estado_final": "BLOQUEADO_POR_CAMBIO_DEL_ARCHIVO_PES",
            "hash_esperado": EXPECTED_PES_HASH,
            "hash_obtenido": pes_hash,
        }
        write_csv(inventory, out_dir / "INVENTARIO_CONGELAMIENTO.csv")
        write_json(summary, out_dir / "RESUMEN_VALIDACION_FINAL.json")
        return 2

    precarga = pd.read_csv(PRECARGA, sep=";", encoding="cp1252", dtype=str, keep_default_na=False).fillna("")
    final = pd.read_csv(PES_FINAL, sep=";", dtype=str, keep_default_na=False).fillna("")
    matriz = pd.read_csv(MATRIZ_VIGENCIA, dtype=str, keep_default_na=False).fillna("")
    audit = pd.read_csv(AUDITORIA_DECISION, dtype=str, keep_default_na=False).fillna("")
    allowed_nac, chile_codes, allowed_res = load_catalogs()
    vig_audit = {trimmed(r["ID_REGISTRO"]): {k: trimmed(v) for k, v in r.items()} for _, r in audit.iterrows()}
    matriz_by_id = {trimmed(r["ID_REGISTRO"]): {k: trimmed(v) for k, v in r.items()} for _, r in matriz.iterrows()}

    final_functional = pd.DataFrame()
    prec_functional = pd.DataFrame()
    for field in FUNCTIONAL_COLUMNS:
        final_functional[field] = final[physical_col(field)]
        prec_functional[field] = precarga[physical_col(field)]
    final_functional.insert(0, "ID_REGISTRO", [f"PRECARGA_{i:04d}" for i in range(1, len(final_functional) + 1)])
    final_functional.insert(1, "FILA_PRECARGA", list(range(1, len(final_functional) + 1)))
    prec_out = precarga.copy()
    prec_out.insert(0, "ID_REGISTRO", [f"PRECARGA_{i:04d}" for i in range(1, len(prec_out) + 1)])
    prec_out.insert(1, "FILA_PRECARGA", list(range(1, len(prec_out) + 1)))

    governance_rows: list[dict[str, Any]] = []
    comparison_rows: list[dict[str, Any]] = []
    difference_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []

    for idx in range(len(final_functional)):
        id_registro = final_functional.loc[idx, "ID_REGISTRO"]
        fila = str(final_functional.loc[idx, "FILA_PRECARGA"])
        matriz_row = matriz_by_id.get(id_registro, {})
        row_final = {field: final_functional.loc[idx, field] for field in FUNCTIONAL_COLUMNS}
        row_comp: dict[str, Any] = {
            "ID_REGISTRO": id_registro,
            "FILA_PRECARGA": fila,
            "CODIGO_IES_NUM": precarga.loc[idx, "CODIGO_IES_NUM"],
            "TIPO_DOCUMENTO": final_functional.loc[idx, "TIPO_DOCUMENTO"],
            "NUM_DOCUMENTO": final_functional.loc[idx, "NUM_DOCUMENTO"],
            "CODIGO_UNICO": final_functional.loc[idx, "CODIGO_UNICO"],
        }
        for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1):
            prec_val = clean(prec_functional.loc[idx, field])
            final_val = clean(final_functional.loc[idx, field])
            norm_prec = normalize_value(field, prec_val)
            norm_final = normalize_value(field, final_val)
            tipo_cambio, origen, fuente, regla, obs_diff = classify_difference(
                field, prec_val, final_val, id_registro, vig_audit, matriz_row
            )
            errors, obs_control = validate_field(field, final_val, row_final, allowed_nac, chile_codes, allowed_res)
            coincide = "SI" if norm_prec == norm_final else "NO"
            modificado = "NO" if coincide == "SI" else "SI"
            gobernado = "SI" if tipo_cambio != "CAMBIO_NO_GOBERNADO" and not any(e.startswith("ERROR_DOMINIO") or e.startswith("ERROR_OBLIGATORIEDAD") or e.startswith("ERROR_NO_SE_PUEDE") or e.startswith("ERROR_PAIS") or e.startswith("ERROR_UNIVERSIDAD") or e.startswith("ERROR_ANIO") or e.startswith("ERROR_FECHA") or e.startswith("ERROR_CODIGO") for e in errors) else "NO"
            resultado = "OK" if gobernado == "SI" and not errors else "ERROR"
            row = {
                "ID_REGISTRO": id_registro,
                "FILA_PRECARGA": fila,
                "CODIGO_IES_NUM": precarga.loc[idx, "CODIGO_IES_NUM"],
                "TIPO_DOCUMENTO": final_functional.loc[idx, "TIPO_DOCUMENTO"],
                "NUM_DOCUMENTO": final_functional.loc[idx, "NUM_DOCUMENTO"],
                "CODIGO_UNICO": final_functional.loc[idx, "CODIGO_UNICO"],
                "POSICION_SIES": pos,
                "CAMPO": field,
                "VALOR_PRECARGA": prec_val,
                "VALOR_FINAL": final_val,
                "VALOR_NORMALIZADO_PRECARGA": norm_prec,
                "VALOR_NORMALIZADO_FINAL": norm_final,
                "COINCIDE_PRECARGA": coincide,
                "FUE_MODIFICADO": modificado,
                "ORIGEN_VALOR_FINAL": origen,
                "FUENTE_RESPALDO": fuente,
                "REGLA_APLICADA": regla,
                "NIVEL_RESPALDO": "ALTO" if origen in {"PRECARGA_OFICIAL", "DECISION_INSTITUCIONAL_DOCUMENTADA", "FUENTE_INSTITUCIONAL"} else ("MEDIO" if origen == "TRANSFORMACION_TECNICA" else "PENDIENTE"),
                "GOBERNADO": gobernado,
                "RESULTADO_CONTROL": resultado,
                "OBSERVACION": " | ".join([obs_diff] + obs_control + errors).strip(" |"),
            }
            governance_rows.append(row)
            row_comp[f"{field}_PRECARGA"] = prec_val
            row_comp[f"{field}_FINAL"] = final_val
            row_comp[f"{field}_IGUAL"] = coincide
            row_comp[f"{field}_TIPO_CAMBIO"] = tipo_cambio
            row_comp[f"{field}_FUENTE"] = fuente
            row_comp[f"{field}_REGLA"] = regla
            row_comp[f"{field}_GOBERNADO"] = gobernado
            if tipo_cambio != "SIN_DIFERENCIA":
                difference_rows.append({
                    "ID_REGISTRO": id_registro,
                    "FILA_PRECARGA": fila,
                    "CAMPO": field,
                    "VALOR_PRECARGA": prec_val,
                    "VALOR_FINAL": final_val,
                    "TIPO_DIFERENCIA": tipo_cambio,
                    "FUENTE_RESPALDO": fuente,
                    "REGLA_APLICADA": regla,
                    "GOBERNADO": gobernado,
                    "OBSERVACION": obs_diff,
                })
            for err in errors:
                error_rows.append({
                    "ID_REGISTRO": id_registro,
                    "FILA_PRECARGA": fila,
                    "CAMPO": field,
                    "VALOR_FINAL": final_val,
                    "ERROR": err,
                    "BLOQUEANTE": "SI",
                })
        comparison_rows.append(row_comp)

    governance = pd.DataFrame(governance_rows)
    comparison = pd.DataFrame(comparison_rows)
    differences = pd.DataFrame(difference_rows)
    errors = pd.DataFrame(error_rows)
    if errors.empty:
        errors = pd.DataFrame(columns=["ID_REGISTRO", "FILA_PRECARGA", "CAMPO", "VALOR_FINAL", "ERROR", "BLOQUEANTE"])
    if differences.empty:
        differences = pd.DataFrame(columns=["ID_REGISTRO", "FILA_PRECARGA", "CAMPO", "VALOR_PRECARGA", "VALOR_FINAL", "TIPO_DIFERENCIA", "FUENTE_RESPALDO", "REGLA_APLICADA", "GOBERNADO", "OBSERVACION"])

    resumen_col_rows = []
    for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1):
        g = governance[governance["CAMPO"].eq(field)]
        resumen_col_rows.append({
            "POSICION": pos,
            "COLUMNA": field,
            "FILAS": len(g),
            "VACIOS": int(g["VALOR_FINAL"].map(trimmed).eq("").sum()),
            "NO_VACIOS": int(g["VALOR_FINAL"].map(trimmed).ne("").sum()),
            "VALORES_UNICOS": int(g["VALOR_FINAL"].nunique()),
            "COINCIDENCIAS_CON_PRECARGA": int(g["COINCIDE_PRECARGA"].eq("SI").sum()),
            "DIFERENCIAS": int(g["COINCIDE_PRECARGA"].eq("NO").sum()),
            "CAMBIOS_GOBERNADOS": int(((g["COINCIDE_PRECARGA"].eq("NO")) & (g["GOBERNADO"].eq("SI"))).sum()),
            "TRANSFORMACIONES_TECNICAS": int(g["REGLA_APLICADA"].eq("TRANSFORMACION_TECNICA_NO_MATERIAL").sum()),
            "CAMBIOS_NO_GOBERNADOS": int(((g["COINCIDE_PRECARGA"].eq("NO")) & (g["GOBERNADO"].eq("NO"))).sum()),
            "ERRORES_DOMINIO": int(g["OBSERVACION"].str.contains("ERROR_DOMINIO|ERROR_NACIONALIDAD_CHILENA|ERROR_FECHA|ERROR_ANIO|ERROR_SEMESTRE", regex=True, na=False).sum()),
            "ERRORES_OBLIGATORIEDAD": int(g["OBSERVACION"].str.contains("ERROR_OBLIGATORIEDAD|ERROR_.*REQUERID", regex=True, na=False).sum()),
            "RESULTADO": "OK" if g["RESULTADO_CONTROL"].eq("OK").all() else "ERROR",
        })
    resumen_col = pd.DataFrame(resumen_col_rows)
    formatos = pd.DataFrame([detect_format(PES_FINAL)])

    duplicate_persona_carrera = int(final_functional.duplicated(subset=["TIPO_DOCUMENTO", "NUM_DOCUMENTO", "CODIGO_UNICO"], keep=False).sum())
    solo_precarga = 0 if len(precarga) == len(final) else max(len(precarga) - len(final), 0)
    solo_final = 0 if len(precarga) == len(final) else max(len(final) - len(precarga), 0)
    cambios_no_gobernados = int(differences["TIPO_DIFERENCIA"].eq("CAMBIO_NO_GOBERNADO").sum())
    errores_bloqueantes = int(errors["BLOQUEANTE"].eq("SI").sum())
    vacios_obligatorios = int(errors["ERROR"].str.contains("ERROR_OBLIGATORIEDAD_VACIO", regex=False, na=False).sum())
    final_header_ok = formatos.loc[0, "encabezado_probable"] == "NO"
    final_delim_ok = formatos.loc[0, "delimitador_probable"] == ","
    can_make_ready = (
        len(final_functional) == 185
        and len(FUNCTIONAL_COLUMNS) == 20
        and len(governance) == 3700
        and cambios_no_gobernados == 0
        and errores_bloqueantes == 0
        and duplicate_persona_carrera == 0
        and solo_precarga == 0
        and solo_final == 0
        and int(final_functional["VIGENCIA"].eq("0").sum()) == 24
        and int(final_functional["VIGENCIA"].eq("1").sum()) == 161
        and int((~final_functional["VIGENCIA"].isin(["0", "1"])).sum()) == 0
        and int(final_functional["CODIGO_UNICO"].map(trimmed).eq("").sum()) == 0
        and vacios_obligatorios == 0
    )

    pes_ready_path = ""
    pes_ready_hash = ""
    desktop_path = ""
    desktop_hash = ""
    desktop_identical = "NO_APLICA"
    if can_make_ready:
        ready_name = f"EXTRANJEROS_REGULARES_2025_PES_READY_{timestamp}.csv"
        ready = out_dir / ready_name
        with ready.open("w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f, delimiter=",", lineterminator="\n", quoting=csv.QUOTE_MINIMAL)
            for _, row in final_functional[FUNCTIONAL_COLUMNS].iterrows():
                writer.writerow([clean(row[col]) for col in FUNCTIONAL_COLUMNS])
        reread = pd.read_csv(ready, header=None, sep=",", dtype=str, keep_default_na=False).fillna("")
        if reread.shape == (185, 20) and reread.iloc[:, 19].eq("0").sum() == 24 and reread.iloc[:, 19].eq("1").sum() == 161:
            desktop = Path.home() / "Desktop" / ready_name
            shutil.copy2(ready, desktop)
            pes_ready_path = str(ready)
            desktop_path = str(desktop)
            pes_ready_hash = sha256(ready)
            desktop_hash = sha256(desktop)
            desktop_identical = "SI" if pes_ready_hash == desktop_hash else "NO"
        else:
            ready.unlink(missing_ok=True)
            can_make_ready = False

    status = "LISTO_PARA_CARGA_PES" if can_make_ready and desktop_identical == "SI" else "BLOQUEADO_POR_VALIDACION_FINAL"
    pruebas_fallidas = []
    if len(governance) != 3700:
        pruebas_fallidas.append("CONTROLES_GOBERNANZA_DISTINTO_3700")
    if cambios_no_gobernados:
        pruebas_fallidas.append("CAMBIOS_NO_GOBERNADOS")
    if errores_bloqueantes:
        pruebas_fallidas.append("ERRORES_BLOQUEANTES")
    if duplicate_persona_carrera:
        pruebas_fallidas.append("DUPLICADOS_PERSONA_CARRERA")
    if vacios_obligatorios:
        pruebas_fallidas.append("CAMPOS_OBLIGATORIOS_VACIOS")
    if not final_header_ok:
        pruebas_fallidas.append("ARCHIVO_REVISADO_TIENE_ENCABEZADO")
    if not final_delim_ok:
        pruebas_fallidas.append("ARCHIVO_REVISADO_NO_USA_COMA")
    if status != "LISTO_PARA_CARGA_PES" and not pruebas_fallidas:
        pruebas_fallidas.append("NO_SE_GENERO_PES_READY")

    resumen = {
        "proceso": "Estudiantes Extranjeros Regulares 2025 - SIES 2026",
        "anio_datos": 2025,
        "estado_final": status,
        "carpeta_validacion": str(out_dir),
        "precarga_utilizada": str(PRECARGA),
        "hash_precarga": sha256(PRECARGA),
        "encoding_precarga": "cp1252",
        "delimitador_precarga": "punto y coma",
        "filas_precarga": int(len(precarga)),
        "columnas_fisicas_precarga": int(len(precarga.columns)),
        "columna_adicional": "CODIGO_IES_NUM",
        "mapeo_pais_origen": "PAIS_ORIGEN -> PAIS_DE_ORIGEN",
        "archivo_final_revisado": str(PES_FINAL),
        "hash_archivo_final_revisado": pes_hash,
        "hash_esperado_archivo_final": EXPECTED_PES_HASH,
        "filas": int(len(final_functional)),
        "columnas": 20,
        "controles_gobernanza_esperados": 3700,
        "controles_gobernanza_ejecutados": int(len(governance)),
        "coincidencias_exactas": int(governance["COINCIDE_PRECARGA"].eq("SI").sum()),
        "transformaciones_tecnicas": int(differences["TIPO_DIFERENCIA"].eq("TRANSFORMACION_TECNICA_NO_MATERIAL").sum()) if not differences.empty else 0,
        "cambios_funcionales_gobernados": int(differences["TIPO_DIFERENCIA"].eq("CAMBIO_FUNCIONAL_GOBERNADO").sum()) if not differences.empty else 0,
        "cambios_no_gobernados": cambios_no_gobernados,
        "errores_bloqueantes": errores_bloqueantes,
        "vacios_obligatorios": vacios_obligatorios,
        "duplicados_persona_carrera": duplicate_persona_carrera,
        "solo_precarga": solo_precarga,
        "solo_final": solo_final,
        "vigencia_0": int(final_functional["VIGENCIA"].eq("0").sum()),
        "vigencia_1": int(final_functional["VIGENCIA"].eq("1").sum()),
        "vigencia_fuera_dominio": int((~final_functional["VIGENCIA"].isin(["0", "1"])).sum()),
        "codigo_unico_vacio": int(final_functional["CODIGO_UNICO"].map(trimmed).eq("").sum()),
        "encabezado_archivo_revisado": formatos.loc[0, "encabezado_probable"],
        "delimitador_archivo_revisado": formatos.loc[0, "delimitador_probable"],
        "codificacion_archivo_revisado": formatos.loc[0, "codificacion_detectada"],
        "bom_archivo_revisado": formatos.loc[0, "bom"],
        "archivo_pes_ready": pes_ready_path,
        "hash_pes_ready": pes_ready_hash,
        "copia_escritorio": desktop_path,
        "hash_copia_escritorio": desktop_hash,
        "copia_escritorio_identica": desktop_identical,
        "pruebas_fallidas": pruebas_fallidas,
    }

    dictionary = pd.DataFrame([
        {
            "POSICION": pos,
            "COLUMNA": field,
            "COLUMNA_FISICA_PRECARGA": physical_col(field),
            "COLUMNA_FISICA_FINAL": physical_col(field),
            "OBLIGATORIEDAD_CONTROLADA": "SI" if field in STRICT_REQUIRED else "CONDICIONAL/O_PERMITIDO_SEGUN_FUENTE",
        }
        for pos, field in enumerate(FUNCTIONAL_COLUMNS, start=1)
    ])
    hashes = pd.DataFrame(inventory_items + [
        {"ITEM": "MATRIZ_GOBERNANZA_185_X_20.csv", "RUTA": str(out_dir / "MATRIZ_GOBERNANZA_185_X_20.csv"), "EXISTE": True, "TAMANO_BYTES": "", "FECHA_MODIFICACION": "", "SHA256": ""},
    ])

    write_csv(inventory, out_dir / "INVENTARIO_CONGELAMIENTO.csv")
    write_csv(governance, out_dir / "MATRIZ_GOBERNANZA_185_X_20.csv")
    write_csv(comparison, out_dir / "COMPARACION_PRECARGA_VS_FINAL_20_COLUMNAS.csv")
    write_csv(differences, out_dir / "DIFERENCIAS_GOBERNANZA_20_COLUMNAS.csv")
    write_csv(errors, out_dir / "ERRORES_VALIDACION_FINAL.csv")
    write_csv(resumen_col, out_dir / "RESUMEN_POR_COLUMNA.csv")
    write_csv(formatos, out_dir / "FORMATOS.csv")
    write_json(resumen, out_dir / "RESUMEN_VALIDACION_FINAL.json")

    report_lines = [
        "# Validacion final gobernanza 20 columnas",
        "",
        f"Estado final: {status}",
        f"Proceso: {resumen['proceso']}",
        f"Archivo final revisado: {PES_FINAL}",
        f"Hash archivo final revisado: {pes_hash}",
        f"Controles ejecutados: {len(governance)}",
        f"Cambios no gobernados: {cambios_no_gobernados}",
        f"Errores bloqueantes: {errores_bloqueantes}",
        f"VIGENCIA 0: {resumen['vigencia_0']}",
        f"VIGENCIA 1: {resumen['vigencia_1']}",
        f"PES READY: {pes_ready_path or 'NO_GENERADO'}",
        "",
        "## Resultado por columna",
        "",
        markdown_table(resumen_col),
    ]
    (out_dir / "REPORTE_VALIDACION_FINAL.md").write_text("\n".join(report_lines) + "\n", encoding="utf-8")

    # Hashes after all main artifacts exist.
    hash_rows = []
    for path in sorted(out_dir.glob("*")):
        if path.is_file() and path.name != "HASHES_VALIDACION_FINAL.sha256":
            hash_rows.append(f"{sha256(path)}  {path}")
    (out_dir / "HASHES_VALIDACION_FINAL.sha256").write_text("\n".join(hash_rows) + "\n", encoding="utf-8")
    hashes = pd.DataFrame([
        {"ARCHIVO": str(path), "SHA256": sha256(path)}
        for path in sorted(out_dir.glob("*"))
        if path.is_file()
    ])

    write_excel(out_dir / "AUDITORIA_GOBERNANZA_20_COLUMNAS_Y_FORMATO_SIES.xlsx", {
        "RESUMEN": pd.DataFrame([resumen]),
        "DICCIONARIO_20_COLUMNAS": dictionary,
        "RESUMEN_POR_COLUMNA": resumen_col,
        "GOBERNANZA_3700": governance,
        "PRECARGA_185": prec_out,
        "FINAL_185": final_functional,
        "COMPARACION_20_COLUMNAS": comparison,
        "DIFERENCIAS": differences,
        "ERRORES": errors,
        "VIGENCIA": matriz,
        "FORMATOS": formatos,
        "HASHES": hashes,
    })

    # Refresh hashes to include the Excel workbook too.
    hash_rows = []
    for path in sorted(out_dir.glob("*")):
        if path.is_file() and path.name != "HASHES_VALIDACION_FINAL.sha256":
            hash_rows.append(f"{sha256(path)}  {path}")
    (out_dir / "HASHES_VALIDACION_FINAL.sha256").write_text("\n".join(hash_rows) + "\n", encoding="utf-8")

    print(json.dumps(resumen, ensure_ascii=False, indent=2))
    return 0 if status == "LISTO_PARA_CARGA_PES" else 2


if __name__ == "__main__":
    raise SystemExit(main_corregido())
