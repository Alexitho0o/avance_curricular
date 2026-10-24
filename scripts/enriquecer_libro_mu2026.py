#!/usr/bin/env python3
"""Enriquece el libro analítico MU2026 mediante un trasplante XML limitado.

El paquete --columnas debe haber sido creado con Artifact Tool. openpyxl se usa
exclusivamente en modo lectura. No ejecuta el motor MU ni modifica sus fuentes.
Contrato: docs/ENRIQUECIMIENTO_CARRERA_SECTOR_MU2026.md.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import copy
import hashlib
import itertools
import json
import os
from pathlib import Path
import posixpath
import re
import sys
import tempfile
from xml.sax.saxutils import escape, quoteattr
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from common.catalogos_sies import cargar_catalogo, normalizar, sha256

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAIN = "06_DETALLE_TOTAL"
SUBSETS = ("04_DETALLE_VIGENTES", "05_DETALLE_NO_VIGENTES")
TARGETS = (MAIN, *SUBSETS)
COLS = {MAIN: ("Q", "R", "S", "T"), **{s: ("Q", "R", "S") for s in SUBSETS}}
WIDTHS = {MAIN: {17: 24, 18: 100, 19: 29, 20: 24}, **{s: {17: 100, 18: 29, 19: 24} for s in SUBSETS}}
HEADERS = ("RUT", "NOMBRE_COMPLETO", "COD_CAR", "COD_SED", "MODALIDAD", "JOR", "VERSION", "OFERTA_SIES", "ANIO_ING_ACT", "SEM_ING_ACT", "ANIO_ING_ORI", "SEM_ING_ORI", "FECHA_MATRICULA", "VIG", "ESTADO_MATRICULA", "ES_VIGENTE")
CELL_RE = re.compile(r'<c\b[^>]*(?:/>|>.*?</c>)', re.S)
ROW_RE = re.compile(r'<row\b[^>]*>.*?</row>', re.S)
FILTER_VALUE = "Sin coincidencia en Hoja1"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def attrs(tag: str) -> dict[str, str]:
    return dict(re.findall(r'([\w:]+)="([^"]*)"', tag.split(">", 1)[0]))


def set_attr(tag: str, key: str, value: object) -> str:
    pattern = rf'\s{re.escape(key)}="[^"]*"'
    replacement = f" {key}={quoteattr(str(value))}"
    if re.search(pattern, tag):
        return re.sub(pattern, replacement, tag, count=1)
    position = tag.find("/>") if tag.endswith("/>") else tag.find(">")
    return tag[:position] + replacement + tag[position:]


def sheet_map(z: ZipFile) -> dict[str, dict]:
    rels = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))}
    result = {}
    for index, s in enumerate(ET.fromstring(z.read("xl/workbook.xml")).find(f"{{{NS}}}sheets")):
        target = rels[s.get(f"{{{REL}}}id")]
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        result[s.get("name")] = {"path": path, "sheet_id": s.get("sheetId"), "index": index}
    return result


def shared_strings(z: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    return ["".join(e.itertext()) if not e.findall(f".//{{{NS}}}t") else "".join(t.text or "" for t in e.findall(f".//{{{NS}}}t")) for e in ET.fromstring(z.read("xl/sharedStrings.xml"))]


def cell_text(cell: ET.Element, strings: list[str]) -> str:
    require(cell.find(f"{{{NS}}}f") is None, "El paquete de columnas debe contener valores, no fórmulas")
    kind = cell.get("t")
    if kind == "s":
        return strings[int(cell.findtext(f"{{{NS}}}v"))]
    if kind == "inlineStr":
        return "".join(t.text or "" for t in cell.findall(f".//{{{NS}}}t"))
    require(kind == "str", "Todas las columnas de enriquecimiento deben contener texto")
    return cell.findtext(f"{{{NS}}}v", "")


def load_mapping(path: Path) -> tuple[dict, dict]:
    document = json.loads(path.read_text(encoding="utf-8"))
    require(document["anio"] == 2026, "El diccionario debe corresponder a 2026")
    catalog = cargar_catalogo(2026)
    by_code = {r["codigo_unico"]: r for r in catalog["oferta"]}
    evidence = document["fuentes"]
    for label in ("catalogo", "publicacion", "sector", "oferta_literal"):
        item = evidence[label]
        require(sha256(Path(item["path"])) == item["sha256"], f"Cambió la evidencia {label}")
    require(json.loads(Path(evidence["catalogo"]["path"]).read_text(encoding="utf-8")) == catalog, "El catálogo documentado no es el catálogo anual activo 2026")
    publication = openpyxl.load_workbook(evidence["publicacion"]["path"], read_only=True, data_only=True)
    iterator = publication[evidence["publicacion"].get("hoja", "Hoja1")].values
    headers = next(iterator)
    positions = {v: i for i, v in enumerate(headers)}
    public_names = {}
    for row in iterator:
        code = row[positions["CÓDIGO CARRERA"]]
        if code in document["por_codigo_unico"]:
            pair = (row[positions["NOMBRE CARRERA"]], row[positions["TIPO DE PLAN DE LA CARRERA"]])
            require(code not in public_names or public_names[code] == pair, "Nombre o plan ambiguo en publicación")
            public_names[code] = pair
    publication.close()
    literal_source = evidence["oferta_literal"]
    literal_sheet = literal_source["hoja"]
    offer = openpyxl.load_workbook(literal_source["path"], read_only=True, data_only=True)
    try:
        iterator = offer[literal_sheet].values
        positions = {v: i for i, v in enumerate(next(iterator))}
        required = ("CODIGO_IES_NUM", "CODIGO_UNICO", "CODIGO_CARRERA", "NOMBRE_CARRERA", "MODALIDAD", "CODIGO_JORNADA", "VERSION_CARRERA")
        require(all(k in positions for k in required), "Oferta literal sin columnas requeridas")
        literal_names = {}
        for number, row in enumerate(iterator, 2):
            code = row[positions["CODIGO_UNICO"]]
            if code is None:
                continue
            require(code not in literal_names, "Código único duplicado en oferta literal")
            name = row[positions["NOMBRE_CARRERA"]]
            require(isinstance(name, str) and bool(name), "Nombre vacío o no textual en oferta literal")
            literal_names[code] = {
                "nombre": name,
                "fila": number,
                "institucion": str(row[positions["CODIGO_IES_NUM"]]),
                "carrera": str(row[positions["CODIGO_CARRERA"]]),
                "modalidad_codigo": str(row[positions["MODALIDAD"]]),
                "jornada": str(row[positions["CODIGO_JORNADA"]]),
                "version": str(row[positions["VERSION_CARRERA"]]),
            }
    finally:
        offer.close()
    for code, item in document["por_codigo_unico"].items():
        require(code in by_code and code in public_names, "Clave sin respaldo 2026")
        cat = by_code[code]
        require(public_names[code] == (item["nombre_oficial"], item["plan"]), "Nombre o plan distinto de la publicación")
        require(normalizar(item["nombre_oficial"]) in {normalizar(n) for n in cat["nombres_fuente"]}, "Nombre ajeno al catálogo")
        require(all(item[k] == cat[k] for k in ("sede", "carrera", "jornada", "version", "modalidad_codigo")), "Dimensiones contradictorias en diccionario")
        require(item["plan_codigo"] == cat["plan"] and item["continuidad"] == (cat["plan"] == "3"), "Continuidad contradictoria")
        require(code in literal_names, "Clave sin nombre en oferta literal 2026")
        literal = literal_names[code]
        require(item["nombre_oferta_literal"] == literal["nombre"], "Nombre distinto del literal de oferta 2026")
        require(item["hoja_oferta_literal"] == literal_sheet and item["fila_oferta_literal"] == literal["fila"], "Ubicación del nombre literal contradictoria")
        require(literal["institucion"] == "162" and all(item[k] == literal[k] for k in ("carrera", "jornada", "version", "modalidad_codigo")), "Dimensiones contradictorias en oferta literal")
        require(code == f"I162S{item['sede']}C{literal['carrera']}J{literal['jornada']}V{literal['version']}", "Identidad de oferta literal contradictoria")
        require(bool(item.get("sector_ipss")), "Sector IPSS sin resolver")
    return document, by_code


def expected_columns(source: Path, mapping: dict) -> tuple[dict, dict, dict]:
    book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    require(all(s in book.sheetnames for s in TARGETS), "Faltan hojas de detalle")
    expected, rowsets, summary, cached_sector = {}, {}, {}, {}
    for name in TARGETS:
        rows = list(book[name].values)
        require(tuple(rows[0][:16]) == HEADERS, "Encabezados A:P distintos del contrato")
        require(all(any(r[:16]) for r in rows[1:]), "Filas vacías dentro del detalle")
        jornada_index = 19 if name == MAIN else 18
        jornada_header = rows[0][jornada_index] if len(rows[0]) > jornada_index else None
        require(jornada_header in (None, "Jornada", "Jornada / modalidad"), "Encabezado de jornada ajeno al contrato")
        jornada_header = jornada_header or "Jornada"
        expected[name] = {1: ("Código Único", "Carrera", "Sector IPSS", jornada_header) if name == MAIN else ("Carrera", "Sector IPSS", jornada_header)}
        counter = Counter()
        sectors, jornadas, codes = Counter(), Counter(), set()
        for number, row in enumerate(rows[1:], 2):
            counter[tuple((type(v).__name__, str(v)) for v in row[:16])] += 1
            cod, sede, modalidad, jornada, version, oferta = map(str, row[2:8])
            code = f"I162S{sede}C{cod}J{jornada}V{version}"
            require(code in mapping, "Clave de estudiante sin correspondencia en catálogo")
            item = mapping[code]
            require((item["sede"], item["carrera"], item["jornada"], item["version"], item["modalidad_codigo"]) == (sede, cod, jornada, version, modalidad), "Dimensiones de fila contradictorias")
            require(oferta == item["clave_oferta"] == f"S{sede}C{cod}M{modalidad}J{jornada}V{version}", "OFERTA_SIES contradictoria")
            career = item["nombre_oferta_literal"]
            values = (career, item["sector_ipss"], item["jornada_legible"])
            expected[name][number] = (code, *values) if name == MAIN else values
            sectors[item["sector_ipss"]] += 1
            jornadas[item["jornada_legible"]] += 1
            codes.add(code)
            if name == MAIN:
                cached_sector[number] = row[18] if len(row) > 18 else None
        rowsets[name] = counter
        summary[name] = {"filas": len(rows) - 1, "claves": len(codes), "sector_ipss": dict(sorted(sectors.items())), "jornada": dict(sorted(jornadas.items()))}
    book.close()
    require(rowsets[SUBSETS[0]] + rowsets[SUBSETS[1]] == rowsets[MAIN], "A:P de vigentes y no vigentes no coincide con total")
    return expected, summary, cached_sector


def artifact_values(path: Path, expected: dict) -> dict:
    result = {}
    with ZipFile(path) as z:
        sheets, strings = sheet_map(z), shared_strings(z)
        require(all(s in sheets for s in TARGETS), "Paquete Artifact sin las tres hojas requeridas")
        for name in TARGETS:
            cells = {}
            for cell in ET.fromstring(z.read(sheets[name]["path"])).findall(f".//{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c"):
                coordinate = cell.get("r")
                require(re.fullmatch(r"[A-Z]+[1-9][0-9]*", coordinate) is not None, "Coordenada inválida")
                col = re.match(r"[A-Z]+", coordinate).group()
                if col not in COLS[name]:
                    # Artifact puede materializar celdas de relleno vacías.
                    require(cell.find(f"{{{NS}}}v") is None and cell.find(f"{{{NS}}}is") is None and cell.find(f"{{{NS}}}f") is None, "Paquete contiene datos fuera de las columnas autorizadas")
                    continue
                require(coordinate not in cells, "Celda duplicada en paquete")
                cells[coordinate] = cell_text(cell, strings)
            wanted = {f"{col}{row}": value for row, values in expected[name].items() for col, value in zip(COLS[name], values)}
            require(cells == wanted, "Contenido de Artifact distinto del diccionario validado")
            result[name] = cells
    return result


def update_widths(xml: str, widths: dict[int, int]) -> str:
    match = re.search(r"<cols>.*?</cols>", xml, re.S)
    require(match is not None, "Hoja sin definición de anchos")
    original = re.findall(r"<col\b[^>]*/>", match.group())
    entries, templates = [], {}
    for tag in original:
        a = attrs(tag); start, end = int(a["min"]), int(a["max"])
        affected = sorted(c for c in widths if start <= c <= end)
        if not affected:
            entries.append((start, tag)); continue
        cursor = start
        for col in affected:
            templates[col] = tag
            if cursor < col:
                entries.append((cursor, set_attr(set_attr(tag, "min", cursor), "max", col - 1)))
            cursor = col + 1
        if cursor <= end:
            entries.append((cursor, set_attr(set_attr(tag, "min", cursor), "max", end)))
    for col, width in widths.items():
        tag = templates.get(col, '<col min="1" max="1"/>')
        for key, value in (("min", col), ("max", col), ("width", width), ("customWidth", 1)):
            tag = set_attr(tag, key, value)
        tag = re.sub(r'\sbestFit="[^"]*"', "", tag)
        entries.append((col, tag))
    replacement = "<cols>" + "".join(t for _, t in sorted(entries)) + "</cols>"
    return xml[:match.start()] + replacement + xml[match.end():]


def patch_sheet(raw: bytes, name: str, values: dict, cached_sector: dict, source_strings: list[str]) -> tuple[bytes, dict]:
    xml = raw.decode("utf-8")
    first_row = ROW_RE.search(xml)
    require(first_row is not None, "Hoja sin encabezado")
    header_cells = {attrs(m.group())["r"] for m in CELL_RE.finditer(first_row.group())}
    missing_widths = {17 + i: WIDTHS[name][17 + i] for i, col in enumerate(COLS[name]) if f"{col}1" not in header_cells}
    row_count = max(int(re.search(r"\d+$", coord).group()) for coord in values)
    hidden_to_show = set()
    cleared = False
    filter_match = re.search(r"<autoFilter\b[^>]*(?:/>|>.*?</autoFilter>)", xml, re.S)
    require(filter_match is not None, "Falta autofiltro existente")
    if name == MAIN and "<filterColumn" in filter_match.group():
        filter_xml = ET.fromstring(filter_match.group().replace(" xr:", " xr_"))
        children = list(filter_xml)
        require(len(children) == 1 and children[0].tag == "filterColumn" and children[0].get("colId") == "18", "Existe otro criterio de filtro; revisión necesaria")
        filters = children[0].findall("filters/filter")
        require(len(filters) == 1 and filters[0].get("val") == FILTER_VALUE, "Filtro sectorial diferente del documentado")
        hidden_to_show = {int(attrs(m.group())["r"]) for m in ROW_RE.finditer(xml) if attrs(m.group()).get("hidden") == "1"}
        require(hidden_to_show == {r for r, val in cached_sector.items() if val != FILTER_VALUE}, "No se puede distinguir filas filtradas de ocultaciones manuales")
        tag = filter_match.group().split(">", 1)[0] + "/>"
        xml = xml[:filter_match.start()] + tag + xml[filter_match.end():]
        xml = re.sub(r"<sheetPr\b[^>]*>", lambda m: re.sub(r'\sfilterMode="[^"]*"', "", m.group()), xml, count=1)
        cleared = True
    touched_rows = set()

    def replace_row(match: re.Match) -> str:
        row = match.group(); number = int(attrs(row)["r"])
        require(1 <= number <= row_count, "Fila fuera del universo esperado")
        touched_rows.add(number)
        cells = {attrs(m.group())["r"]: m.group() for m in CELL_RE.finditer(row)}
        for col in COLS[name]:
            coordinate = f"{col}{number}"
            if coordinate in cells:
                existing = ET.fromstring(f'<root xmlns="{NS}">{cells[coordinate]}</root>')[0]
                if existing.find(f"{{{NS}}}f") is None and existing.get("t") in {"s", "inlineStr", "str"}:
                    if cell_text(existing, source_strings) == values[coordinate]:
                        continue
            style_source = cells.get(f"R{number}") if name == MAIN and col == "Q" else cells.get(coordinate)
            if style_source is None and name == MAIN:
                style_source = cells.get(f"S{number}")
            if style_source is None:
                style_source = cells.get(f"P{number}")
            require(style_source is not None, "No existe celda fuente de estilo")
            style = attrs(style_source).get("s")
            new = f'<c r="{coordinate}"' + (f' s="{style}"' if style else "") + f' t="inlineStr"><is><t xml:space="preserve">{escape(values[coordinate])}</t></is></c>'
            if coordinate in cells:
                row = row.replace(cells[coordinate], new, 1)
            else:
                row = row.replace("</row>", new + "</row>")
        opening, rest = row.split(">", 1)
        opening += ">"
        opening = set_attr(opening, "spans", f"1:{16 + len(COLS[name])}")
        if number in hidden_to_show:
            opening = re.sub(r'\shidden="1"', "", opening)
        return opening + rest

    xml = ROW_RE.sub(replace_row, xml)
    require(touched_rows == set(range(1, row_count + 1)), "Faltan filas de destino")
    end_col = COLS[name][-1]
    xml = re.sub(r"<dimension\b[^>]*/>", lambda m: set_attr(m.group(), "ref", f"A1:{end_col}{row_count}"), xml, count=1)
    xml = re.sub(r"<autoFilter\b[^>]*>", lambda m: set_attr(m.group(), "ref", f"A1:{end_col}{row_count}"), xml, count=1)
    if missing_widths:
        xml = update_widths(xml, missing_widths)
    ET.fromstring(xml)
    return xml.encode("utf-8"), {"filtro_sector_limpiado": cleared, "filas_desocultadas_por_filtro": len(hidden_to_show)}


def patch_defined_names(raw: bytes, sheets: dict, expected: dict) -> bytes:
    xml = raw.decode("utf-8")
    by_index = {str(sheets[name]["index"]): name for name in TARGETS}
    def replace(match: re.Match) -> str:
        text = match.group(); a = attrs(text)
        if a.get("name") != "_xlnm._FilterDatabase" or a.get("localSheetId") not in by_index:
            return text
        name = by_index[a["localSheetId"]]
        require(re.search(r"!\$A\$1:\$[P-T]\$\d+<", text) is not None, "Rango de filtro inesperado")
        return re.sub(r"!\$A\$1:\$[P-T]\$\d+<", f"!$A$1:${COLS[name][-1]}${len(expected[name])}<", text)
    return re.sub(r"<definedName\b[^>]*>.*?</definedName>", replace, xml, flags=re.S).encode("utf-8")


def patch_calc_chain(raw: bytes, main_id: str) -> tuple[bytes, int]:
    current_id = None
    removed = 0
    def replace(match: re.Match) -> str:
        nonlocal current_id, removed
        tag = match.group(); a = attrs(tag)
        current_id = a.get("i", current_id)
        require(current_id is not None, "calcChain sin sheetId inicial")
        if current_id == main_id and re.fullmatch(r"[QRS][1-9][0-9]*", a.get("r", "")):
            removed += 1
            return ""
        require("i" in a, "calcChain usa sheetId implícito fuera del alcance; revisar antes de editar")
        return tag
    xml = re.sub(r"<c\b[^>]*/>", replace, raw.decode("utf-8"))
    return xml.encode("utf-8"), removed


def verify(source: Path, result: Path, replacements: dict, expected: dict, values: dict) -> dict:
    unchanged_cells = 0
    with ZipFile(source) as original, ZipFile(result) as final:
        require(original.namelist() == final.namelist(), "Cambió la estructura del paquete ZIP")
        for name in original.namelist():
            require(final.read(name) == replacements.get(name, original.read(name)), "Parte ZIP alterada fuera del trasplante")
        sheets = sheet_map(original)
        for name in TARGETS:
            before = {attrs(m.group())["r"]: m.group() for m in CELL_RE.finditer(original.read(sheets[name]["path"]).decode())}
            after = {attrs(m.group())["r"]: m.group() for m in CELL_RE.finditer(final.read(sheets[name]["path"]).decode())}
            affected = set(values[name])
            require({k: v for k, v in before.items() if k not in affected} == {k: v for k, v in after.items() if k not in affected}, "Se modificaron celdas XML ajenas al alcance")
    old = openpyxl.load_workbook(source, read_only=True, data_only=False)
    new = openpyxl.load_workbook(result, read_only=True, data_only=False)
    require(old.sheetnames == new.sheetnames, "Cambió el orden o nombre de hojas")
    for name in old.sheetnames:
        left, right = old[name], new[name]
        require(left.max_row == right.max_row, "Cambió el número de filas")
        for number, (row_a, row_b) in enumerate(zip(left.iter_rows(), right.iter_rows()), 1):
            for column, (cell_a, cell_b) in enumerate(itertools.zip_longest(row_a, row_b), 1):
                if name in TARGETS and 17 <= column <= 16 + len(COLS[name]):
                    require(cell_b is not None and cell_b.value == expected[name][number][column - 17] and cell_b.data_type == "s", "Valor final de enriquecimiento incorrecto")
                    continue
                signature = lambda c: (c.value, c.data_type, getattr(c, "_style_id", 0)) if c is not None else (None, "n", 0)
                require(signature(cell_a) == signature(cell_b), "Cambió valor, fórmula o estilo fuera del alcance")
                unchanged_cells += 1
    old.close(); new.close()
    return {"celdas_ajenas_verificadas": unchanged_cells, "xml_celdas_ajenas_identico": True, "valores_formulas_estilos_ajenos_identicos": True, "partes_no_afectadas_identicas": True}


def install_atomic(temporary: Path, destination: Path) -> None:
    if destination.exists():
        require(sha256(temporary) == sha256(destination), "El destino ya existe con contenido diferente")
        temporary.unlink()
    else:
        # Enlace exclusivo evita sobrescribir un archivo creado entre preflight y commit.
        os.link(temporary, destination)
        temporary.unlink()


def run(args: argparse.Namespace) -> dict:
    for key in ("entrada", "catalogo_nombres", "columnas", "salida", "auditoria"):
        setattr(args, key, getattr(args, key).expanduser().resolve())
    require(args.entrada != args.salida, "La salida debe ser distinta de la entrada")
    require(args.auditoria not in {args.entrada, args.salida, args.catalogo_nombres, args.columnas}, "La auditoría debe tener un destino independiente")
    if args.salida.exists():
        require(not os.path.samefile(args.entrada, args.salida), "Entrada y salida son el mismo archivo")
    document, _ = load_mapping(args.catalogo_nombres)
    expected, summary, cached_sector = expected_columns(args.entrada, document["por_codigo_unico"])
    values = artifact_values(args.columnas, expected)
    replacements, visibility = {}, {}
    removed = 0
    with ZipFile(args.entrada) as source:
        sheets = sheet_map(source)
        strings = shared_strings(source)
        for name in TARGETS:
            path = sheets[name]["path"]
            replacements[path], visibility[name] = patch_sheet(source.read(path), name, values[name], cached_sector, strings)
        replacements["xl/workbook.xml"] = patch_defined_names(source.read("xl/workbook.xml"), sheets, expected)
        if "xl/calcChain.xml" in source.namelist():
            replacements["xl/calcChain.xml"], removed = patch_calc_chain(source.read("xl/calcChain.xml"), sheets[MAIN]["sheet_id"])
        args.salida.parent.mkdir(parents=True, exist_ok=True)
        handle, temp_name = tempfile.mkstemp(prefix=".mu2026_", suffix=".xlsx", dir=args.salida.parent)
        os.close(handle)
        temporary = Path(temp_name)
        try:
            with ZipFile(temporary, "w") as target:
                target.comment = source.comment
                for info in source.infolist():
                    target.writestr(copy(info), replacements.get(info.filename, source.read(info.filename)))
            checks = verify(args.entrada, temporary, replacements, expected, values)
            audit = {"schema_version": 1, "anio": 2026, "estado": "VALIDADO", "entrada": str(args.entrada), "salida": str(args.salida), "hashes": {"entrada_sha256": sha256(args.entrada), "salida_sha256": sha256(temporary), "diccionario_sha256": sha256(args.catalogo_nombres), "columnas_artifact_sha256": sha256(args.columnas)}, "fuentes": document["fuentes"], "hojas": summary, "control": {"multiset_A_P_vigentes_mas_no_vigentes_igual_total": True, "cobertura_nombre_sector_completa": True, "claves_modalidad_oferta_concordantes": True, **checks}, "visibilidad": visibility, "calcChain_formulas_retiradas": removed, "partes_modificadas": sorted(name for name, content in replacements.items() if content != source.read(name))}
            raw_audit = (json.dumps(audit, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode()
            if args.auditoria.exists():
                require(args.auditoria.read_bytes() == raw_audit, "Auditoría existente diferente")
            args.auditoria.parent.mkdir(parents=True, exist_ok=True)
            audit_fd, audit_name = tempfile.mkstemp(prefix=".mu2026_audit_", suffix=".json", dir=args.auditoria.parent)
            audit_temp = Path(audit_name)
            try:
                with os.fdopen(audit_fd, "wb") as stream:
                    stream.write(raw_audit)
                    stream.flush(); os.fsync(stream.fileno())
                install_atomic(temporary, args.salida)
                install_atomic(audit_temp, args.auditoria)
            finally:
                audit_temp.unlink(missing_ok=True)
        finally:
            temporary.unlink(missing_ok=True)
    return audit


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("entrada", "catalogo-nombres"):
        parser.add_argument(f"--{name}", type=Path, required=True)
    for name in ("columnas", "salida", "auditoria", "preparar-columnas"):
        parser.add_argument(f"--{name}", type=Path)
    args = parser.parse_args()
    try:
        if args.preparar_columnas:
            document, _ = load_mapping(args.catalogo_nombres)
            expected, _, _ = expected_columns(args.entrada, document["por_codigo_unico"])
            payload = {"sheets": [{"name": name, "start": "Q", "end": COLS[name][-1], "matrix": [list(v) for _, v in sorted(expected[name].items())]} for name in TARGETS]}
            destination = args.preparar_columnas.expanduser().resolve()
            require(destination not in {args.entrada.resolve(), args.catalogo_nombres.resolve()}, "Destino de payload coincide con fuente")
            raw = (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode()
            if destination.exists():
                require(destination.read_bytes() == raw, "Payload existente diferente")
            else:
                destination.parent.mkdir(parents=True, exist_ok=True)
                with destination.open("xb") as stream:
                    stream.write(raw)
            print(json.dumps({"estado": "PREPARADO", "payload": str(destination), "sin_datos_personales": True}, ensure_ascii=False))
            return 0
        require(all((args.columnas, args.salida, args.auditoria)), "Indicar --columnas, --salida y --auditoria para enriquecer")
        audit = run(args)
    except (ValueError, KeyError, OSError, ET.ParseError) as error:
        print(f"ERROR: {type(error).__name__}: {error}", file=sys.stderr)
        return 2
    print(json.dumps({"estado": audit["estado"], "salida": audit["salida"], "filas": {k: v["filas"] for k, v in audit["hojas"].items()}, "partes_modificadas": audit["partes_modificadas"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
