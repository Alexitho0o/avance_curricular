"""Diccionario y controles agregados UMAS; datos individuales sólo en memoria/privado."""
import collections,csv,datetime,hashlib,itertools,json,re,subprocess,sys,unicodedata
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from common.diccionario_umas import texto,clave,rut_validado,cabeceras_unicas,etiqueta_inscripcion,periodo_umas,comparacion_univoca,sensibilidad,semantica
OUT=ROOT/'data/diccionario_umas';PRIVATE=Path('/Users/alexi/.local/share/ipss/diccionario_umas_20261028')
def sha(p):
    with open(p,'rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def read(p):return json.loads(Path(p).read_text())
def write(p,x):Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n')
def nonempty(v):return v is not None and texto(v).upper() not in ('','NULL','NONE','NAN')
def normalize(v):return texto(v).upper().strip()
def family(name):
    n=name.upper()
    return 'DATOS_ALUMNOS' if 'DATOSDEALUMNOS' in n else 'PROMEDIOS' if 'PROMEDIOSDEALUMNOS' in n else 'PERMANENCIA' if 'PERMANENCIA' in n else 'GESTIONSGE'
def header_rows(sheet):
    iterator=sheet.iter_rows(values_only=True);first=list(itertools.islice(iterator,12))
    core={'CODCLI','CODCARPR','CODCARR','RUT','RUTALUMNO','ESTADOACADEMICO','OBSERVACION','JORNADA','ANOMATRICULA'}
    scores=[len({clave(v) for v in row}&core) for row in first]
    if not scores or max(scores)<3:return None,[],iter(())
    i=max(range(len(scores)),key=lambda i:scores[i])
    return i+1,list(first[i]),itertools.chain(first[i+1:],iterator)
def num(v):
    try:return float(texto(v).replace(',','.'))
    except (ValueError,TypeError):return None
def pick(row,*names):
    for n in names:
        if n in row and nonempty(row[n]):return row[n]
    return None
def attributes(row,perma=False):
    label=periodo_umas(pick(row,'ANOPERIODOINGRESO')) if perma else None
    return {'carrera':normalize(pick(row,'CARRERAALUM' if perma else 'CODCARPR','CODCARR')),'carrera_actual':normalize(pick(row,'CODIGOCARRERAACTUAL' if perma else 'CODIGOCARRERA')),'jornada':normalize(pick(row,'JORNADA')),'sede':normalize(pick(row,'CODSEDE' if perma else 'SEDE')),'ingreso_anio':texto(pick(row,'ANOINGRESO')),'ingreso_periodo':label[1] if label else texto(pick(row,'PERIODOINGRESO')),'estado':normalize(pick(row,'ESTADOALUMNO' if perma else 'ESTADOACADEMICO'))}

def main():
    OUT.mkdir(parents=True,exist_ok=True);PRIVATE.mkdir(parents=True,exist_ok=True,mode=0o700)
    assert not Path('/Users/alexi/Downloads/~$Permanencia.xlsx').exists(),'Excel abierto: detener'
    cat=read(ROOT/'data/catalogo_fuentes/CATALOGO_FUENTES.json')
    sources=[r for r in cat['catalogo'] if any(k in r['archivo'].upper() for k in ('DATOSDEALUMNOS','PROMEDIOSDEALUMNOS','GESTIONSGE','PERMANENCIA'))]
    contract={clave(r['columna']):r for r in read(ROOT/'data/reglas_negocio/20261021_datos_alumnos/contrato_columnas.json')['columnas']}
    perms={};da_indexes={};profiles=[];secrets=set();prom=collections.defaultdict(set);prom_equiv=collections.defaultdict(set)
    aliases=collections.defaultdict(set);header_changes=[];errors=[]
    for source in sources:
        path=Path(source['ruta_consolidada']);fam=family(path.name)
        assert sha(path)==source['sha256'],'Fuente alterada: '+path.name
        if fam=='PERMANENCIA':assert not path.with_name('~$'+path.name).exists(),'Excel abierto: detener'
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        meta={'archivo':path.name,'familia':fam,'ruta':str(path),'sha256':source['sha256'],'bytes':source['bytes'],'hojas':[],'clase':'B metadatos/observaciones; C lectura y perfiles; E semántica sin respaldo'}
        if fam=='DATOS_ALUMNOS':da_indexes[path.name]=collections.defaultdict(lambda:collections.defaultdict(set))
        for sheet in wb:
            hr,original,iterator=header_rows(sheet)
            sm={'hoja':sheet.title,'filas_declaradas':sheet.max_row,'columnas_declaradas':sheet.max_column,'fila_cabecera':hr,'filas_leidas':0,'columnas':[]}
            if hr is None:
                sm['estado']='E_HOJA_AUXILIAR_CABECERA_NO_ACREDITADA';sm['columnas']=[{'posicion':i,'nombre_normalizado':'COLUMNA_'+str(i).zfill(3),'sensibilidad':'excluida','tipo':'no perfilado','significado':'Cabecera no acreditada: no interpretar filas como encabezados ni publicar sus valores.','clase':'E'} for i in range(1,sheet.max_column+1)]
                meta['hojas'].append(sm);continue
            unique=cabeceras_unicas(original,fam=='PERMANENCIA');stats=[]
            for i,(old,new) in enumerate(zip(original,unique),1):
                sensitive=sensibilidad(new,fam);meaning,cl,grain=semantica(new,fam,contract)
                item={'posicion':i,'nombre_original':texto(old),'nombre_normalizado':new,'sensibilidad':sensitive,'significado':meaning,'clase':cl,'grano_variable':grain,'tipos_observados':set(),'valores_observados':collections.Counter(),'vacios':0,'n':0,'min':None,'max':None}
                stats.append(item);aliases[clave(old)].add((fam,path.name,sheet.title,new))
                if texto(old)!=new and fam=='PERMANENCIA':header_changes.append({'posicion':i,'original':old,'normalizado':new,'clase':'C alias de lectura; libro original intacto'})
            sheet_cids=set();sheet_years=collections.Counter();dup=0
            for raw in iterator:
                if not any(nonempty(v) for v in raw):continue
                sm['filas_leidas']+=1
                row={clave(k):v for k,v in zip(unique,raw)}
                for s,v in zip(stats,raw):
                    if not nonempty(v):s['vacios']+=1;continue
                    s['tipos_observados'].add(type(v).__name__);s['n']+=1
                    if s['sensibilidad']=='analítica':
                        s['valores_observados'][texto(v)]+=1;n=num(v)
                        if n is not None:s['min']=n if s['min'] is None else min(n,s['min']);s['max']=n if s['max'] is None else max(n,s['max'])
                    elif any(t in clave(s['nombre_original']) for t in ('CODCLI','RUT','MAIL','FONO','APELLIDO','PATERNO','MATERNO')) or (clave(s['nombre_original']) in ('NOMBRE','NOMBRES','NOMBREALUMNO') and fam!='PERMANENCIA'):
                        val=texto(v)
                        if len(val)>=6:secrets.add(val)
                cid=texto(pick(row,'CODCLI'))
                if cid:
                    if cid in sheet_cids:dup+=1
                    sheet_cids.add(cid)
                y=texto(pick(row,'ANO','ANOMATRICULA','ANIO'))
                if y:sheet_years[y]+=1
                if fam=='DATOS_ALUMNOS' and cid:
                    for k,v in attributes(row).items():
                        if nonempty(v):da_indexes[path.name][cid][k].add(v)
                elif fam=='PERMANENCIA' and cid:
                    assert cid not in perms,'CODCLI no único en Permanencia; detener unión'
                    perms[cid]=row
                elif fam=='PROMEDIOS':
                    career=texto(pick(row,'CODCARR','CODCARPR'));plan=texto(pick(row,'PLANDEESTUDIO','CODPESTUD'));ramo=texto(pick(row,'CODRAMO'));eq=texto(pick(row,'RAMOEQUIV'))
                    if career and plan and ramo:prom[career,plan].add(ramo)
                    if career and plan and eq:prom_equiv[career,plan].add(eq)
            for s in stats:
                s['tipos_observados']=sorted(s['tipos_observados'])
                if s['sensibilidad']!='analítica':
                    for key in ('vacios','n','min','max'):s.pop(key)
                    s['valores_observados']={'estado':'NO_EXPORTADOS_PRIVACIDAD_O_EXCLUSION'}
                else:
                    cnt=s['valores_observados'];s['distintos']=len(cnt)
                    s['valores_observados']={'frecuencias':dict(sorted(cnt.items()))} if len(cnt)<=150 else {'cardinalidad':len(cnt),'nota':'Dominio largo; sin ejemplos individuales.'}
                sm['columnas'].append(s)
            sm.update(estado='B_C_PERFILADO',codcli_distintos=len(sheet_cids),filas_codcli_repetido=dup,anos_observados=dict(sheet_years),grano='asignatura/inscripción CODCLI-año-período-ramo-sección' if fam=='PROMEDIOS' else 'evento de gestión; no posee CODCLI acreditado' if fam=='GESTIONSGE' else 'expediente CODCLI con posibles repeticiones en hojas derivadas')
            meta['hojas'].append(sm)
        wb.close();assert sha(path)==source['sha256'];profiles.append(meta)
        print(json.dumps({'perfilado':path.name,'hojas':len(meta['hojas'])},ensure_ascii=False),flush=True)
    assert len(perms)==4982
    # Controles propios de Permanencia: no se serializan identificadores ni gestiones.
    rut=collections.Counter();types_rut=collections.Counter();flags=collections.defaultdict(collections.Counter);displaced=collections.Counter();unknown_labels=0;counts_ano=collections.Counter();listlen=collections.Counter();pending_lists=[];year_discrepancy=0
    header_periods={clave('Inscribe_asignaturas_en_t'+str(p)):p for p in (1,2,3)}
    for cid,row in perms.items():
        r=rut_validado(row['RUT'],formato='cuerpo');rut[r['estado']]+=1;types_rut[r['formato']]+=1;counts_ano[texto(row['ANO'])]+=1
        seen=collections.defaultdict(set)
        for field,header_p in header_periods.items():
            tag=etiqueta_inscripcion(row[field])
            if tag:
                y,p,yes=tag;seen[y,p].add(yes)
                if p!=header_p:displaced['columna_t'+str(header_p)]+=1
            else:unknown_labels+=1
        for p in (1,2,3):
            states=seen.get((2026,p),set());label='sin_etiqueta' if not states else 'contradictorio' if len(states)>1 else 'si_tomo' if next(iter(states)) else 'no_tomo'
            flags['2026-'+str(p)][label]+=1
        ep=periodo_umas(row['ANOPERIODOINGRESO'])
        if ep and ep[0]!=texto(row['ANOINGRESO']):year_discrepancy+=1
        s='' if row.get('LISTARAMOSPEND') is None else str(row['LISTARAMOSPEND']);listlen[len(s)]+=1;pending_lists.append(s)
    controls={'clase':'B observado; C control; E semántica/cobertura pendiente','permanencia':{'fecha_extraccion':'2026-10-27','filas':4982,'codcli_distintos':4982,'campo_ano':dict(counts_ano),'sin_columna_plan':True,'fotografia_no_historia':True},'rut':{'validacion_modulo11':dict(rut),'formatos':dict(types_rut),'nota':'Pegar DV sólo se reconoce como candidato de formato y se valida; no se corrigen fuentes ni se inventan ceros iniciales.'},'inscripciones_por_etiqueta':{p:dict(c) for p,c in sorted(flags.items())},'etiquetas_desplazadas_respecto_cabecera':dict(displaced),'celdas_sin_etiqueta_valida':unknown_labels,'ingreso_anio_vs_etiqueta_discrepancias':year_discrepancy,'lista_ramos_pend':{'longitud_maxima':max(listlen),'filas_en_maximo':listlen[max(listlen)],'longitudes_mayores':[{'caracteres':n,'filas':c} for n,c in sorted(listlen.items(),reverse=True)[:12]],'dictamen':'C: concentración en 800 caracteres compatible con límite/truncamiento del reporte. E: no se conoce la consulta exportadora; no tratar lista como completa.','valores_no_exportados':True},'columnas_duplicadas_renombradas':header_changes,'fuente_original_no_editada':True}
    # Cada variable tiene su propio denominador comparable; ambigüedades no se fuerzan.
    agreement=[]
    for filename,index in da_indexes.items():
        common=set(perms)&set(index);variable={k:collections.Counter() for k in attributes(next(iter(perms.values())),True)}
        for cid in common:
            target=attributes(perms[cid],True)
            for k in variable:variable[k][comparacion_univoca(index[cid].get(k,set()),target[k])]+=1
        detail={}
        for k,cnt in variable.items():
            n=cnt['concordantes']+cnt['diferentes'];detail[k]={**{s:cnt[s] for s in ('concordantes','diferentes','sin_dato','ambiguos')},'comparables':n,'tasa_concordancia_pct':100*cnt['concordantes']/n if n else None}
            assert sum(cnt.values())==len(common)
        agreement.append({'fuente_datos_alumnos':filename,'codcli_permanencia':4982,'codcli_comunes':len(common),'sin_codcli_comun':4982-len(common),'variables':detail,'clase':'B/C concordancia del estado de dos extractos, sin vigencia histórica; no se elige el más nuevo por nombre de archivo'})
    controls['concordancia_por_codcli']=agreement
    # Mallas disponibles: fuentes H112 y 5810 congeladas; no ejecutar motores anteriores.
    m=read(ROOT/'outputs/avance_historico_fase3_20261027/MANIFEST.json')['fuentes'];malla_sources={k:m[k] for k in ('listado','h112a','h112b_excel','h112b_csv','manual')}
    for s in malla_sources.values():assert sha(s['ruta'])==s['sha256']
    w=openpyxl.load_workbook(m['listado']['ruta'],read_only=True,data_only=True);s=w['BBDD Bruta'];it=s.iter_rows(values_only=True);h=next(it);raw_mallas=collections.defaultdict(set);plan_careers=collections.defaultdict(set);course_names=collections.defaultdict(set)
    for values in it:
        r={clave(k):v for k,v in zip(h,values)};c=texto(r['CODCARR']);p=texto(r['CODPESTUD']);ramo=texto(r['CODRAMO'])
        if c and p and ramo:raw_mallas[c,p].add(ramo);plan_careers[p].add(c)
        if ramo:course_names[ramo].add(clave(r.get('NOMBRE')))
    w.close()
    w=openpyxl.load_workbook(m['h112a']['ruta'],read_only=True,data_only=True);it=w['05_FUENTE_DETALLE_ASIGNATURAS'].iter_rows(values_only=True);h=next(it);governed=collections.defaultdict(set)
    for values in it:
        r=dict(zip(h,values));key=(texto(r['CODIGO_UNICO']),texto(r['CODPESTUD']));u=texto(r.get('UNIDAD_ID_USADA'))
        if u:governed[key].add(u)
    w.close()
    w=openpyxl.load_workbook(m['h112b_excel']['ruta'],read_only=True,data_only=True);it=w['03_CARGA_SIES_16769_21C'].iter_rows(values_only=True);h=next(it);loads=[dict(zip(h,row)) for row in it];it=w['02_H112A_HOJA10'].iter_rows(values_only=True);h=next(it);plan_by_code={texto(r[0]):dict(zip(h,r)) for r in it};w.close()
    carga_by_career=collections.defaultdict(list)
    for r in loads:
        assert texto(r['TIPO_UNIDAD_MEDIDA'])=='1'
        code=texto(r['CODIGO_UNICO']);plan=texto(plan_by_code[code]['CODPESTUD'])
        # Carrera UMAS desde plan literal del listado, sin alias por nombre ni prefijo.
        for c in plan_careers[plan]:carga_by_career[c].append({'codigo_sies':code,'plan_umas':plan,'total_carga_5810':num(r['TOTAL_UNIDADES_MEDIDA']),'unidades_detalle_gobernado':len(governed[code,plan]),'unidad':'asignaturas/cursos/módulos (TIPO_UNIDAD_MEDIDA=1)','clase':'B/C referencia2026 datos2025; no acredita plan individual de Permanencia'})
    per_career=collections.defaultdict(collections.Counter)
    for row in perms.values():per_career[texto(row['CARRERAALUM'])][texto(row['TOTALRAMOSMALLA'])]+=1
    check=[]
    for c,cnt in sorted(per_career.items()):
        totals={num(v) for v in cnt if num(v) is not None};known_raw=[{'plan':p,'ramos_distintos':len(rs)} for (cc,p),rs in raw_mallas.items() if cc==c]
        load=carga_by_career[c];lt={x['total_carga_5810'] for x in load};rt={r['ramos_distintos'] for r in known_raw};n=sum(cnt.values())
        prom_counts=[{'plan':p,'codramo_distintos_observados':len(rs),'ramoequiv_distintos_observados':len(prom_equiv[cc,p])} for (cc,p),rs in prom.items() if cc==c]
        check.append({'carrera_umas':c,'expedientes':n,'total_ramos_malla_observado':dict(cnt),'es_constante_entre_no_vacios':len(totals)==1,'totales_no_numericos_o_vacios':sum(v for k,v in cnt.items() if num(k) is None),'mallas_listado_por_plan':known_raw,'referencias_carga_5810':load,'PROMEDIOS_por_plan':prom_counts,'con_total_coincidente_alguna_5810':sum(n for v,n in cnt.items() if num(v) in lt),'con_total_distinto_5810':sum(n for v,n in cnt.items() if num(v) is not None and num(v) not in lt) if lt else None,'sin_referencia_5810':not bool(lt),'con_total_coincidente_alguna_malla_listado':sum(n for v,n in cnt.items() if num(v) in rt),'estado':'E_ATRIBUCION_DE_PLAN: sólo comparación numérica por carrera; no se asigna plan ni se corrige total'})
    # Evidencia adicional del límite de800: nombre final parcial contra catálogo de ramos.
    prefix_cut=0
    for s in pending_lists:
        if len(s)!=800:continue
        last=s.split(' / ')[-1];match=re.match(r'^([A-Za-z0-9_]+)-(.+)',last)
        if match:
            code,name=match.groups();name=clave(name)
            if len(name)>=5 and any(full.startswith(name) and len(full)>len(name) for full in course_names.get(code,set())):prefix_cut+=1
    controls['lista_ramos_pend']['filas_800_con_nombre_final_prefijo_incompleto_de_catalogo']=prefix_cut
    # Las fuentes de04 no están materializadas; consultar manifiestos por Git en memoria.
    old='3067c976b19101eca1cb0912f14ea70525c96d34';historical_path='avance_curricular_2026/04_gobernanza_mallas/03_auditorias/AUDITORIA_MALLAS_COMPARTIDAS_END_TO_END_20260701_090314/03_RESULTADOS/09_MANIFIESTO_TABLA_CANONICA_LOGISTICA.json'
    blob=subprocess.check_output(['git','show',old+':'+historical_path],cwd=ROOT);hist= json.loads(blob)
    mallas={'clase':'A definición unidad; B totales observados; C comparación; E plan no informado','referencia_proceso':'Avance Curricular2026 informa2025; Permanencia es foto27-10-2026. No se presupone vigencia de todos los planes en2026.','fuentes':malla_sources,'gobernanza_04':{'ruta':str(ROOT/'avance_curricular_2026/04_gobernanza_mallas'),'archivos_locales_materializados':sum(p.is_file() for p in (ROOT/'avance_curricular_2026/04_gobernanza_mallas').rglob('*')),'manifest_recuperado_de_Git':{'commit':old,'ruta':historical_path,'sha256':hashlib.sha256(blob).hexdigest()},'maestro_tsv_citado':hist['fuente_canonica'],'maestro_tsv_sha256_citado':hist['hash_fuente_canonica'],'payload_maestro_tsv_disponible':Path(hist['fuente_canonica']).exists(),'alternativa_verificada':'H112a detalle gobernado, carga5810H112b y listadoRE+CO; no se restituye ni modifica el árbol histórico'},'por_carrera':check,'resumen':{'carreras':len(check),'expedientes':sum(x['expedientes'] for x in check),'carreras_total_constante':sum(x['es_constante_entre_no_vacios'] for x in check),'carreras_varios_totales':sum(not x['es_constante_entre_no_vacios'] and len([v for v in x['total_ramos_malla_observado'] if num(v) is not None])>1 for x in check),'coincidencias_numericas_5810':sum(x['con_total_coincidente_alguna_5810'] for x in check),'diferencias_numericas_5810':sum(x['con_total_distinto_5810'] for x in check if x['con_total_distinto_5810'] is not None),'sin_referencia_5810':sum(x['expedientes'] for x in check if x['sin_referencia_5810'])},'limites':['Permanencia no trae plan: igualdad de total no prueba igualdad de malla.','Listado bruto puede repetir ramos por requisitos; se cuentan códigos distintos, sin asumir obligatoriedad de todas las filas.','PROMEDIOS es actividad observada y no acredita malla completa; convalidaciones/equivalencias no se fusionan automáticamente.','HistorialGit prueba manifiesto; TSV canónico ausente no se declara reconstruido.']}
    assert mallas['resumen']['expedientes']==4982
    equivalents={'clase':'B/C variantes físicas; E equivalencias semánticas condicionadas','reglas':[{'variable':'expediente','variantes':['CODCLI'],'regla':'texto literal; no extraer ingreso desde patrón; no aparece en GESTIONSGE1703','clase':'B/C'},{'variable':'documento','variantes':['RUT','RUT_ALUMNO','DIG','DV'],'regla':'conservar original y formato; validarDV por módulo11; no fabricarDV ni cruzar por nombre','clase':'C'},{'variable':'carrera_expediente','variantes':['CODCARPR','CODCARR','CARRERA_ALUM'],'regla':'candidatas a misma dimensión, verificar por CODCLI y corte; CODIGO_CARRERA_ACTUAL y CODIGOCARRERA se conservan separados','clase':'B/C/E'},{'variable':'jornada','variantes':['JORNADA'],'regla':'letras UMAS preservadas; no inferir modalidad ni código SIES','clase':'B/C'},{'variable':'sede','variantes':['SEDE','COD_SEDE'],'regla':'comparar código UMAS literal; NOMBRE en Permanencia es sede, en DatosAlumnos es persona','clase':'B/C'},{'variable':'ingreso','variantes':['ANOINGRESO','ANO_INGRESO','PERIODOINGRESO','ANO_PERIODO_INGRESO'],'regla':'separar AAAA-P; no renombrar períodos4/5; no asumir cohorte institucional','clase':'B/C'},{'variable':'estado_estudiante','variantes':['ESTADOACADEMICO','ESTADO_ACADEMICO','ESTADO_ALUMNO'],'regla':'foto del extracto, no historial; ESTADO en PROMEDIOS es otro campo: asignatura','clase':'B/C'},{'variable':'nivel','variantes':['NIVEL'],'regla':'NIVEL de PROMEDIOS corresponde a asignatura; en otras fuentes es administrativo. No homologar a avance completado','clase':'B/C'},{'variable':'inscripcion_periodo','variantes':list(header_periods),'regla':'parsear etiqueta del valor; nombre de columna no determina período','clase':'B/C'}],'renombres_permanencia':header_changes,'apariciones_fisicas':[{'columna_canonica_lexica':k,'apariciones':[{'familia':f,'archivo':a,'hoja':s,'columna':c} for f,a,s,c in sorted(v)]} for k,v in sorted(aliases.items())]}
    dic={'estado':'DICCIONARIO_EXPLORATORIO_GOBERNADO_NO_CAMBIA_CARGAS','fecha':'2026-10-28','alcance':'Todas las extracciones UMAS de CATALOGO_FUENTES, deduplicadas por SHA-256; todas sus hojas. Sin diccionarios paralelos por edición.','fuentes':profiles,'columnas_privadas':'Sólo metadatos y tipos; sin valores, frecuencias, ejemplos de contacto, morosidad o comentarios.','precedencia':'Contrato20261021 y regla de proceso/corte; ninguna frecuencia observada constituye regla oficial.','referencias':[{'ruta':str(ROOT/p),'sha256':sha(ROOT/p)} for p in ('data/reglas_negocio/20261021_datos_alumnos/contrato_columnas.json','docs/SEGUIMIENTO_RUT_Y_CONVERSION_PERIODOS_20261023.md')],'controles':controls}
    write(OUT/'DICCIONARIO_UMAS.json',dic);write(OUT/'EQUIVALENCIAS_UMAS.json',equivalents);write(OUT/'COMPROBACION_MALLAS.json',mallas)
    write(OUT/'CONTROLES_UMAS.json',controls)
    # Tokens sólo privados para escaneo de salidas, sin exportar relaciones individuales.
    secretfile=PRIVATE/'TOKENS_PRIVACIDAD.json';write(secretfile,sorted(secrets));secretfile.chmod(0o600)
    lines=['# Diccionario UMAS · corte28-10-2026','','B/C/E: diccionario de esquemas y observaciones; no modifica cargas, bases, denominadores ni tasas. A sólo en las definiciones oficiales citadas.','','## Permanencia movida [D/B/C]','','Movimiento autorizado a13_UMAS_PERMANENCIA_RESTRINGIDO/20261027/Permanencia.xlsx; SHA-256 antes/después56437ff17e5efd9c0f2e398e8f8ba4490730a43c970e44676357925553198c6b.4.982 filas y CODCLI distintos. Fecha27-10-2026; ANO2026 en4.981 filas y2025 en una, conservada. Foto actual sin columna de plan.','','## Cobertura de esquemas','','| Fuente | Familia | Hojas | Filas perfiladas |','|---|---|---:|---:|']
    for p in profiles:lines.append(f"| {p['archivo']} | {p['familia']} | {len(p['hojas'])} | {sum(s['filas_leidas'] for s in p['hojas'])} |")
    lines+=['','Inventario completo de columnas, posición original/normalizada, tipos, significado, clase, sensibilidad y valores agregados: DICCIONARIO_UMAS.json. Hojas auxiliares sin cabecera acreditada se excluyen de interpretación y no publican valores.','','## Controles de lectura [B/C/E]','', 'RUT: '+json.dumps(controls['rut'],ensure_ascii=False),'','| Período real de etiqueta | Sí tomó | No tomó | Sin etiqueta | Contradictorio |','|---|---:|---:|---:|---:|']
    for p,c in sorted(flags.items()):lines.append(f"| {p} | {c['si_tomo']} | {c['no_tomo']} | {c['sin_etiqueta']} | {c['contradictorio']} |")
    lines+=['',f"Tres cabeceras duplicadas se renombran en la lectura, sin editar el XLSX: {len(header_changes)} columnas originales/derivadas. El sufijo distingue ACADEMIA/FINANZAS.",'',f"LISTA_RAMOS_PEND: máximo{max(listlen)} caracteres; {listlen[max(listlen)]} filas en ese máximo. {prefix_cut} filas de800 terminan en prefijo incompleto de un nombre del catálogo. C: evidencia de corte; E: consulta exportadora y lista completa no acreditadas. No usar lista como total de malla.",'','## Concordancia por CODCLI [B/C]','','Tasa por variable: concordantes/(concordantes+diferentes). Vacíos y ambigüedades se muestran aparte; no se elige primera fila. Son estados de extractos, no vigencia histórica.','','| Datos Alumnos | CODCLI comunes | Carrera % | Jornada % | Sede % | Ingreso año % | Estado % |','|---|---:|---:|---:|---:|---:|---:|']
    for a in agreement:
        percentages=[a['variables'][k]['tasa_concordancia_pct'] for k in ('carrera','jornada','sede','ingreso_anio','estado')];lines.append('| '+a['fuente_datos_alumnos']+' | '+str(a['codcli_comunes'])+' | '+' | '.join('Pendiente' if v is None else f'{v:.2f}%' for v in percentages)+' |')
    lines+=['','## Comprobación de mallas por carrera [A/B/C/E]','','A: Instructivo AvanceCurricular2026 p.9, TIPO_UNIDAD_MEDIDA=1 significa asignaturas/cursos/módulos. Se contrasta carga5810 de43 programas, detalleH112, listadoRE+CO y ramos observados en las cuatro extracciones PROMEDIOS.','','| Carrera | Expedientes | Totales distintos en Permanencia | ¿Constante? | Coincide algún total5810 | Difiere de5810 | Sin referencia5810 |','|---|---:|---|---|---:|---:|---|']
    for r in check:lines.append(f"| {r['carrera_umas']} | {r['expedientes']} | {', '.join(r['total_ramos_malla_observado'])} | {'Sí' if r['es_constante_entre_no_vacios'] else 'No'} | {r['con_total_coincidente_alguna_5810']} | {r['con_total_distinto_5810'] if r['con_total_distinto_5810'] is not None else 'No aplica'} | {'Sí' if r['sin_referencia_5810'] else 'No'} |")
    lines+=['','No se atribuye plan a Permanencia. Coincidencia numérica no acredita malla; diferencias por plan/fecha quedan E. Los directorios04 están vacíos localmente; se consultó su manifiesto histórico en Git sin restaurar el árbol. El payloadTSV canónico no está disponible: se conservan referencias verificadas H112, listado y5810 como contraste.','','## Privacidad y reproducción','','Contactos, nombres de personas, documentos, morosidad y comentarios quedan fuera de los valores y distribuciones públicas. No se publican cruces individuales. GESTIONSGE no aporta CODCLI acreditado y no se fuerza esa equivalencia.','', 'Reproducir: `python3 scripts/construir_diccionario_umas.py`. No mover nuevamente Permanencia; el script lee la ubicación consolidada y verifica hash. No ejecuta motores anuales ni cargas.']
    (OUT/'DICCIONARIO_UMAS.md').write_text('\n'.join(lines)+'\n')
    from scripts.finalizar_diccionario_umas import main as finalizar
    finalizar()
    print(json.dumps({'fuentes':len(profiles),'codcli_permanencia':len(perms),'periodos':{k:dict(v) for k,v in flags.items()},'mallas':mallas['resumen'],'sin_commit_push':True},ensure_ascii=False))
if __name__=='__main__':main()
