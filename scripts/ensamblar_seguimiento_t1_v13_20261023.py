"""Sustituye exclusivamente las hojas 52–54 y valida fórmulas, agregados y conservación."""
import sys,re,json,zipfile,hashlib,shutil,subprocess,xml.etree.ElementTree as E
from pathlib import Path
from xml.sax.saxutils import escape
import openpyxl
O=Path(sys.argv[1]).resolve();R=Path(__file__).resolve().parents[1]
meta=json.loads((O/'MANIFEST_SEGUIMIENTO.json').read_text());src=Path(meta['origen_excel'])
dst=O/'RETENCION_IPSS_SEGUIMIENTO_T1_T2_T3_20261023_V13.xlsx'
if dst.exists():raise RuntimeError('La versión final ya existe; no sobrescribir')
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
for p,h in meta['fuentes'].items():assert sha(p)==h,p
z=zipfile.ZipFile(src);new=zipfile.ZipFile(O/'SUPLEMENTO_QA.xlsx');NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda x:'{'+NS+'}'+x;E.register_namespace('',NS)
strings=[''.join(n.itertext()) for n in E.fromstring(new.read('xl/sharedStrings.xml'))]
base=z.read('xl/styles.xml').decode();b=E.fromstring(base);a=E.fromstring(new.read('xl/styles.xml'));offset={};num_map={}
for tag in ['numFmts','fonts','fills','borders','cellXfs']:
 old=b.find(q(tag));extra=a.find(q(tag));offset[tag]=len(old) if old is not None else 0
 if extra is None:continue
 if tag=='numFmts':
  nextid=max([163]+[int(n.get('numFmtId')) for n in old] if old is not None else [163])+1
  for n in extra:num_map[int(n.get('numFmtId'))]=nextid;n.set('numFmtId',str(nextid));nextid+=1
 if tag=='cellXfs':
  for n in extra:
   for attr,t in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:n.set(attr,str(int(n.get(attr,'0'))+offset[t]))
   n.set('numFmtId',str(num_map.get(int(n.get('numFmtId','0')),int(n.get('numFmtId','0')))));n.set('xfId','0')
 children=''.join(E.tostring(n,encoding='unicode').replace(' xmlns="'+NS+'"','') for n in extra)
 found=re.search(rf'<{tag}\b[^>]*>.*?</{tag}>',base,re.S)
 if found:
  chunk=re.sub(r'count="\d+"',f'count="{offset[tag]+len(extra)}"',found.group(),count=1).replace(f'</{tag}>',children+f'</{tag}>');base=base[:found.start()]+chunk+base[found.end():]
 else:base=base.replace('<fonts',f'<{tag} count="{len(extra)}">{children}</{tag}><fonts',1)
def clean(txt):
 txt=txt.lstrip('\ufeff').replace('<x:','<').replace('</x:','</').replace('xmlns:x=','xmlns=')
 def cell(m):
  st=m.group()
  if re.search(r'\bt="s"',st):
   idx=int(re.search(r'<v>(\d+)</v>',st).group(1));st=re.sub(r'\bt="s"','t="inlineStr"',st);st=re.sub(r'<v>.*?</v>',lambda m:'<is><t xml:space="preserve">'+escape(strings[idx])+'</t></is>',st,flags=re.S)
  if re.search(r'\bs="\d+"',st):st=re.sub(r'\bs="(\d+)"',lambda m:f's="{int(m.group(1))+offset["cellXfs"]}"',st,count=1)
  else:st=st.replace('<c ','<c s="'+str(offset['cellXfs'])+'" ',1)
  return st
 return re.sub(r'<c\b[^>]*(?<!/)>.*?</c>',cell,txt,flags=re.S)
names=['52_Codigos_SIES_Revision','53_Codigos_T1_Confirmados','54_Contextos_y_Pendientes'];repl={'xl/styles.xml':base}
for i,j in [(1,52),(2,53),(3,54)]:
 repl[f'xl/worksheets/sheet{j}.xml']=clean(new.read(f'xl/worksheets/sheet{i}.xml').decode())
 rp=f'xl/worksheets/_rels/sheet{i}.xml.rels'
 if rp in new.namelist():
  rr=E.fromstring(new.read(rp));assert len(rr)==1
  for rel in rr:
   assert rel.get('Type').endswith('/table');target=rel.get('Target').lstrip('/');dest=f'xl/tables/table{j-19}.xml';tt=E.fromstring(new.read(target));tt.set('id',str(j-19));repl[dest]=E.tostring(tt,encoding='unicode');rel.set('Target','/'+dest)
  repl[f'xl/worksheets/_rels/sheet{j}.xml.rels']=E.tostring(rr,encoding='unicode')
assert set(repl).issubset(z.namelist())
with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as out:
 for n in z.namelist():out.writestr(n,repl[n].encode() if n in repl else z.read(n))
zz=zipfile.ZipFile(dst);assert zz.testzip() is None;assert len(zz.namelist())==len(set(zz.namelist()))
assert set(zz.namelist())==set(z.namelist())
changed=[n for n in z.namelist() if z.read(n)!=zz.read(n)];assert set(changed)<=set(repl)
for n in zz.namelist():
 if n.endswith(('.xml','.rels')):E.fromstring(zz.read(n))
for i in range(1,52):assert zz.read(f'xl/worksheets/sheet{i}.xml')==z.read(f'xl/worksheets/sheet{i}.xml')
charts=[n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n)];assert len(charts)==13
assert all(zz.read(n)==z.read(n) for n in charts)
w=openpyxl.load_workbook(dst,read_only=True,data_only=True);wf=openpyxl.load_workbook(dst,read_only=True,data_only=False);assert len(w.sheetnames)==54
forms=json.loads((O/'FORMULAS.json').read_text());assert len(forms)==302
t=json.loads((O/'TABLAS_SEGUIMIENTO.json').read_text());checks=0;expected={}
for r,x in enumerate(t['origen_periodo'],6):expected[(names[1],f'I{r}')]=(x[6]/x[7] if x[7]>0 else 'No aplica') if isinstance(x[6],(int,float)) else 'Pendiente'
for r in range(6,12):
 y=2024+(r-6)//2;p=2+(r-6)%2;sel=[x for x in t['origen_periodo'] if x[0]==y and x[2]==p];cs=[x for x in t['contextos'] if x[0]==y and x[2]==p]
 vals={'D':sum(x[5] for x in sel),'E':600 if y==2025 and p==2 else sum(x[6] for x in sel),'F':sum(x[7] for x in sel)};vals['G']=vals['E']/vals['F']
 for col,v in vals.items():expected[(names[0],f'{col}{r}')]=v
 rr=r+11;vals2={'A':y,'B':2025 if y==2024 else y,'C':p,'D':sum(x[12] for x in sel),'E':sum(x[10] for x in sel),'F':sum(x[13] for x in cs if x[8]=='B/C Documental anual'),'G':sum(x[13] for x in cs if x[8]=='D Interno aprobado'),'H':sum(x[13] for x in cs if x[8]=='E Pendiente'),'I':sum(x[13] for x in cs)}
 for col,v in vals2.items():expected[(names[0],f'{col}{rr}')]=v
for col in 'DEFGHI':expected[(names[0],f'{col}24')]=sum(expected[(names[0],f'{col}{r}')] for r in range(17,23))
assert len(expected)==len(forms)
for x in forms:
 assert wf[x['sheet']][x['cell']].value==x['formula'];got=w[x['sheet']][x['cell']].value;want=expected[(x['sheet'],x['cell'])]
 assert (abs(got-want)<1e-12 if isinstance(want,(int,float)) and isinstance(got,(int,float)) else got==want),(x,got,want);checks+=1
for y,r,source_row in [(2024,6,7),(2025,8,8),(2026,10,11)]:
 for row,col in [(r,'F'),(r+1,'H')]:assert abs(w[names[0]][f'G{row}'].value-w['18_Foto_Intertrimestral'][f'{col}{source_row}'].value)<1e-12;checks+=1
errors=[]
for name in names:
 for row in w[name]:
  for cell in row:
   if cell.data_type=='e':errors.append([name,cell.coordinate,cell.value])
assert not errors
ids=set();tn=set()
for n in zz.namelist():
 if not re.fullmatch(r'xl/tables/table\d+\.xml',n):continue
 tb=E.fromstring(zz.read(n));assert tb.get('id') not in ids and tb.get('name') not in tn;ids.add(tb.get('id'));tn.add(tb.get('name'))
 if n not in ['xl/tables/table34.xml','xl/tables/table35.xml']:assert zz.read(n)==z.read(n);continue
 cols=tb.find(q('tableColumns'));assert int(cols.get('count'))==len(cols);assert tb.find(q('autoFilter')).get('ref')==tb.get('ref')
 name=names[1 if n.endswith('34.xml') else 2];assert [w[name].cell(5,k+1).value for k in range(len(cols))]==[x.get('name') for x in cols]
assert len(ids)==35
orig=json.loads((Path.home()/'.local/share/ipss/aplicacion_diccionario_intertrimestral_20261023_125556/ORIGENES_T1.json').read_text());private=json.loads((Path.home()/'.local/share/ipss/destinos_calendario_v7_20261023/CONTEXTOS_POR_ORIGEN_PERIODO.json').read_text())
personal={str(x[k]) for x in orig for k in ('persona','documento_original') if x.get(k)}|{x['codcli'] for x in private}
for name in names:
 for row in w[name]:
  for cell in row:assert str(cell.value) not in personal
for p,h in meta['fuentes'].items():assert sha(p)==h,p
git={'rama':subprocess.check_output(['git','branch','--show-current'],cwd=R,text=True).strip(),'HEAD':subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip()};assert git==meta['git_antes']
desktop=Path.home()/'Desktop'/dst.name;assert not desktop.exists();shutil.copy2(dst,desktop);assert sha(desktop)==sha(dst)
qa={'estado':'OK_CON_PENDIENTES_E','hojas_originales_conservadas_bytes':51,'graficos_conservados_bytes':13,'tablas_previas_conservadas_bytes':33,'tablas_actualizadas_validas':2,'formulas_recalculadas_y_contrastadas':len(forms),'controles_independientes':checks,'errores_calculo_hojas_actualizadas':errors,'partes_modificadas':changed,'fuentes_sin_cambio':True,'sin_identificadores_en_hojas_actualizadas':True,'denominadores_originales':[558,1150,1471],'denominador_heredado_T3_2026':1467,'nuevos_cierres':0,'limites':['Recálculo con Artifact Tool; no se validó apertura en Excel nativo.','T2-2025 conserva600/1150; desglose pendiente frente a reproducción611.','No certifica las tasas proxy como vigencia histórica al cierre.','104 contextos pendientes no equivalen a104 personas ni a deserción.'],'git_antes':meta['git_antes'],'git_despues':git,'commit':False,'push':False,'archivo':str(dst),'sha256':sha(dst),'escritorio':str(desktop),'sha256_escritorio':sha(desktop)}
(O/'QA_EXCEL.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)+'\n')
meta.update({'resultado':qa,'artefactos_agregados':{p.name:sha(p) for p in O.glob('*.json') if p.name!='MANIFEST_SEGUIMIENTO.json'},'scripts':{str(p.relative_to(R)):sha(p) for p in (R/'scripts').glob('*v13*20261023*') if p.is_file()}})
(O/'MANIFEST_SEGUIMIENTO.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n');print(json.dumps(qa,ensure_ascii=False))
