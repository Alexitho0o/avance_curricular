"""Exportaciones de reconstruccion Fase 2A."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from . import VERSION
from .configuracion import EXPECTED_UNIVERSES, project_root
from .utilidades import sha256_file


def json_default(value: Any) -> Any:
    """Serializa objetos no nativos."""

    if isinstance(value, Path):
        return str(value)
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return str(value)


def audit_dir() -> Path:
    """Crea auditoria nueva sin sobrescribir."""

    target = project_root() / "auditorias" / f"kpi_rotacion_reconstruccion_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    target.mkdir(parents=True, exist_ok=False)
    return target


def write_csv(path: Path, frame: pd.DataFrame) -> Path:
    """Escribe CSV con separador estable."""

    frame.to_csv(path, index=False, sep=";", encoding="utf-8-sig")
    return path


def records(frame: pd.DataFrame) -> list[dict[str, object]]:
    """Convierte frame a registros JSON."""

    if frame.empty:
        return []
    return frame.where(pd.notna(frame), "").to_dict(orient="records")


def format_workbook(path: Path) -> None:
    """Aplica formato minimo al Excel ejecutivo."""

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = fill
    workbook.save(path)


def safe_summary(
    kpi: pd.DataFrame,
    reentries: pd.DataFrame,
    mobility: pd.DataFrame,
    validations: pd.DataFrame,
    sources: pd.DataFrame,
) -> dict[str, object]:
    """Construye resumen sin datos personales."""

    return {
        "version": VERSION,
        "universos_aprobados": EXPECTED_UNIVERSES,
        "fuentes": records(sources[["anio", "archivo", "sha256", "estado"]]),
        "kpi_por_par": records(kpi),
        "reingresos_total": int(len(reentries)),
        "movilidad_total": int(len(mobility)),
        "validaciones": validations.groupby("SEVERIDAD").size().to_dict() if not validations.empty else {},
        "estado_publicacion": "NO_PUBLICABLE_FASE_2A",
    }


def markdown_summary(payload: dict[str, object]) -> str:
    """Genera resumen Markdown sin datos personales."""

    lines = [
        "# Reconstruccion Rotacion Docente SIES Fase 2A",
        "",
        "Resultado reproducible desde fuentes gobernadas locales. No se consultaron fuentes externas ni se integro al flujo PES.",
        "",
        "## KPI Por Par",
        "",
    ]
    for row in payload["kpi_por_par"]:
        lines.append(
            f"- {row['ANIO_BASE']}-{row['ANIO_COMPARACION']}: base={row['DOTACION_BASE']} siguiente={row['DOTACION_SIGUIENTE']} rotan={row['ROTAN']} tasa={row['TASA_OFICIAL_ROTACION']}"
        )
    lines.extend(
        [
            "",
            f"- Reingresos detectados: {payload['reingresos_total']}",
            f"- Movilidades registradas: {payload['movilidad_total']}",
            "- Estado: NO_PUBLICABLE_FASE_2A",
        ]
    )
    return "\n".join(lines) + "\n"


def export_excel(
    path: Path,
    summary: pd.DataFrame,
    panel: pd.DataFrame,
    permanecen: pd.DataFrame,
    rotan: pd.DataFrame,
    nuevos: pd.DataFrame,
    reentries: pd.DataFrame,
    mobility: pd.DataFrame,
    desaggs: dict[str, pd.DataFrame],
    validations: pd.DataFrame,
    audit: pd.DataFrame,
) -> Path:
    """Genera Excel ejecutivo solicitado."""

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="RESUMEN", index=False)
        panel.to_excel(writer, sheet_name="PANEL", index=False)
        permanecen.to_excel(writer, sheet_name="PERMANECEN", index=False)
        rotan.to_excel(writer, sheet_name="ROTAN", index=False)
        nuevos.to_excel(writer, sheet_name="NUEVOS", index=False)
        reentries.to_excel(writer, sheet_name="REINGRESOS", index=False)
        mobility.to_excel(writer, sheet_name="MOVILIDAD", index=False)
        desaggs["sexo"].to_excel(writer, sheet_name="SEXO", index=False)
        desaggs["formacion"].to_excel(writer, sheet_name="FORMACIÓN", index=False)
        desaggs["horas"].to_excel(writer, sheet_name="HORAS", index=False)
        desaggs["programa"].to_excel(writer, sheet_name="PROGRAMAS", index=False)
        desaggs["jerarquia"].to_excel(writer, sheet_name="JERARQUIA", index=False)
        desaggs["adscripcion"].to_excel(writer, sheet_name="ADSCRIPCION", index=False)
        desaggs["tipo_contractual"].to_excel(writer, sheet_name="TIPO_CONTRACTUAL", index=False)
        validations.to_excel(writer, sheet_name="VALIDACIONES", index=False)
        audit.to_excel(writer, sheet_name="AUDITORÍA", index=False)
    format_workbook(path)
    return path


def export_all(
    longitudinal: pd.DataFrame,
    panel: pd.DataFrame,
    detail: pd.DataFrame,
    kpi: pd.DataFrame,
    reentries: pd.DataFrame,
    mobility: pd.DataFrame,
    desaggs: dict[str, pd.DataFrame],
    validations: pd.DataFrame,
    sources: pd.DataFrame,
    commands: list[str],
) -> dict[str, Path]:
    """Exporta todos los artefactos Fase 2A."""

    target = audit_dir()
    outputs: dict[str, Path] = {}
    outputs["longitudinal"] = write_csv(target / "01_base_longitudinal.csv", longitudinal)
    outputs["panel"] = write_csv(target / "02_panel_presencia.csv", panel)
    outputs["permanecen"] = write_csv(target / "03_permanecen.csv", detail[detail["ESTADO_ROTACION"].eq("PERMANECE")])
    outputs["rotan"] = write_csv(target / "04_rotan.csv", detail[detail["ESTADO_ROTACION"].eq("ROTA")])
    outputs["nuevos"] = write_csv(target / "05_nuevos.csv", detail[detail["ESTADO_ROTACION"].eq("NUEVO INGRESO")])
    outputs["reingresos"] = write_csv(target / "06_reingresos.csv", reentries)
    outputs["movilidad"] = write_csv(target / "07_movilidad.csv", mobility)
    outputs["kpi"] = write_csv(target / "08_kpi_por_par.csv", kpi)
    outputs["desglose_sexo"] = write_csv(target / "09_desglose_sexo.csv", desaggs["sexo"])
    outputs["desglose_formacion"] = write_csv(target / "10_desglose_formacion.csv", desaggs["formacion"])
    outputs["desglose_horas"] = write_csv(target / "11_desglose_horas.csv", desaggs["horas"])
    outputs["desglose_cargo"] = write_csv(target / "12_desglose_cargo.csv", desaggs["cargo"])
    outputs["desglose_programa"] = write_csv(target / "13_desglose_programa.csv", desaggs["programa"])
    outputs["validaciones"] = write_csv(target / "14_validaciones.csv", validations)
    payload = safe_summary(kpi, reentries, mobility, validations, sources)
    outputs["resumen_json"] = target / "15_resumen.json"
    outputs["resumen_json"].write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=json_default) + "\n", encoding="utf-8")
    outputs["resumen_md"] = target / "16_resumen.md"
    outputs["resumen_md"].write_text(markdown_summary(payload), encoding="utf-8")
    audit = pd.DataFrame(
        [
            {"CAMPO": "version", "VALOR": VERSION},
            {"CAMPO": "commands", "VALOR": " | ".join(commands)},
            {"CAMPO": "estado_publicacion", "VALOR": "NO_PUBLICABLE_FASE_2A"},
            {"CAMPO": "oneDrive_usado", "VALOR": "NO"},
        ]
    )
    outputs["excel"] = export_excel(
        target / "KPI_ROTACION_DOCENTE_RECONSTRUIDO.xlsx",
        kpi,
        panel,
        detail[detail["ESTADO_ROTACION"].eq("PERMANECE")],
        detail[detail["ESTADO_ROTACION"].eq("ROTA")],
        detail[detail["ESTADO_ROTACION"].eq("NUEVO INGRESO")],
        reentries,
        mobility,
        desaggs,
        validations,
        audit,
    )
    manifest = {
        "version": VERSION,
        "audit_dir": str(target),
        "commands": commands,
        "artifacts": {key: {"path": str(path), "sha256": sha256_file(path)} for key, path in outputs.items() if path.exists()},
        "fuentes": records(sources[["anio", "archivo", "sha256", "estado", "ruta_gobernada", "ruta_normalizada"]]),
        "estado_publicacion": "NO_PUBLICABLE_FASE_2A",
    }
    outputs["manifest"] = target / "17_manifest.json"
    outputs["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=json_default) + "\n", encoding="utf-8")
    outputs["desglose_jerarquia"] = write_csv(target / "18_desglose_jerarquia.csv", desaggs["jerarquia"])
    outputs["desglose_adscripcion"] = write_csv(target / "19_desglose_adscripcion.csv", desaggs["adscripcion"])
    outputs["desglose_tipo_contractual"] = write_csv(target / "20_desglose_tipo_contractual.csv", desaggs["tipo_contractual"])
    outputs["audit_dir"] = target
    return outputs
