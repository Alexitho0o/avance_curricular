#!/usr/bin/env python3
"""
Generador oficial del Sistema Institucional de Indicadores
de Personal Académico — SIIPA.

Principios obligatorios
-----------------------
1. Leer exclusivamente fuentes normalizadas gobernadas.
2. Detectar correctamente archivos CSV delimitados por punto y coma.
3. No modificar archivos fuente.
4. Calcular todos los KPI analíticos en Python.
5. Usar Excel sólo como capa de presentación.
6. No depender de referencias estructuradas masivas.
7. No depender de recálculo de Microsoft Excel.
8. Preservar los KPI oficiales congelados de rotación.
9. Mantener RAW, cálculos, KPI, parámetros y validaciones separados.
10. No ejecutar acciones Git.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import sys
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[2]

PROJECT = ROOT / "personal_academico_2026"

NORMALIZED_DIR = (
    PROJECT
    / "data"
    / "normalized_restricted"
)

DEFAULT_OUTPUT = (
    Path.home()
    / "Desktop"
    / "PERSONAL_ACADEMICO_SIIPA_KPI_v4.0.xlsx"
)

EXPECTED_UNIVERSES = {
    2022: 71,
    2023: 70,
    2024: 85,
    2025: 124,
    2026: 202,
}

EXPECTED_ROTATION = {
    (2022, 2023): (71, 70, 36, 35, 34, 49.295775),
    (2023, 2024): (70, 85, 39, 31, 46, 44.285714),
    (2024, 2025): (85, 124, 50, 35, 74, 41.176471),
    (2025, 2026): (124, 202, 102, 22, 100, 17.741935),
}

SIIPA_SHEET_NAMES = [
    "00_PORTADA",
    "01_RESUMEN_EJECUTIVO",
    "02_KPI_MAESTRO",
    "03_ESTABILIDAD",
    "04_DIVERSIDAD",
    "05_CAPACIDAD_HORAS",
    "06_FORMACION_CARGO",
    "07_CALIDAD_DATOS",
    "08_DATOS_CRUZADOS",
    "09_RAW_2022",
    "10_RAW_2023",
    "11_RAW_2024",
    "12_RAW_2025",
    "13_RAW_2026",
    "90_DICCIONARIO_KPI",
    "91_PARAMETROS",
    "92_CALCULOS",
    "93_VALIDACIONES",
]


@dataclass(frozen=True)
class AnnualSource:
    year: int
    path: Path
    encoding: str
    delimiter: str
    headers: list[str]
    rows: list[dict[str, str]]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)

    return digest.hexdigest()


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalized_name(value: Any) -> str:
    value = text(value).upper()

    value = unicodedata.normalize("NFD", value)
    value = "".join(
        character
        for character in value
        if unicodedata.category(character) != "Mn"
    )

    value = re.sub(r"[^A-Z0-9]+", "_", value)

    return value.strip("_")


def normalized_business_text(value: Any) -> str:
    value = text(value).upper()
    value = unicodedata.normalize("NFD", value)

    value = "".join(
        character
        for character in value
        if unicodedata.category(character) != "Mn"
    )

    value = re.sub(r"\s+", " ", value)

    return value.strip()


def detect_encoding(data: bytes) -> tuple[str, str]:
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ):
        try:
            return encoding, data.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise RuntimeError("No fue posible determinar la codificación.")


def valid_header(fields: list[str]) -> bool:
    normalized = {
        normalized_name(field)
        for field in fields
    }

    return {
        "TIPO_DOCUMENTO",
        "NUM_DOCUMENTO",
        "DV",
    }.issubset(normalized)


def detect_delimiter(decoded: str) -> str:
    lines = [
        line
        for line in decoded.splitlines()
        if line.strip()
    ]

    if not lines:
        raise RuntimeError("La fuente está vacía.")

    evaluations = []

    for delimiter in (";", ",", "\t", "|"):
        fields = next(
            csv.reader(
                [lines[0]],
                delimiter=delimiter,
            )
        )

        evaluations.append(
            (
                valid_header(fields),
                len(fields),
                delimiter,
            )
        )

    valid = [
        result
        for result in evaluations
        if result[0]
    ]

    if not valid:
        raise RuntimeError(
            "Ningún delimitador produjo una cabecera documental válida."
        )

    return max(valid, key=lambda result: result[1])[2]


def read_source(year: int) -> AnnualSource:
    path = (
        NORMALIZED_DIR
        / f"personal_academico_{year}.csv"
    )

    if not path.exists():
        raise FileNotFoundError(path)

    raw = path.read_bytes()
    encoding, decoded = detect_encoding(raw)
    delimiter = detect_delimiter(decoded)

    with path.open(
        "r",
        encoding=encoding,
        newline="",
    ) as file:
        reader = csv.DictReader(
            file,
            delimiter=delimiter,
        )

        rows = list(reader)
        headers = reader.fieldnames or []

    if not valid_header(headers):
        raise RuntimeError(
            f"Cabecera inválida para {year}: {headers[:5]}"
        )

    return AnnualSource(
        year=year,
        path=path,
        encoding=encoding,
        delimiter=delimiter,
        headers=headers,
        rows=rows,
    )


def find_column(
    headers: list[str],
    aliases: list[str],
) -> str | None:
    index = {
        normalized_name(header): header
        for header in headers
    }

    for alias in aliases:
        column = index.get(normalized_name(alias))

        if column:
            return column

    return None


def canonical_key(
    row: dict[str, str],
    columns: dict[str, str | None],
) -> str | None:
    document_type = text(
        row.get(columns["tipo_documento"])
    ).upper()

    document_number = re.sub(
        r"[^0-9A-Z]",
        "",
        text(
            row.get(columns["numero_documento"])
        ).upper(),
    )

    dv = re.sub(
        r"[^0-9K]",
        "",
        text(row.get(columns["dv"])).upper(),
    )

    if not document_type or not document_number or not dv:
        return None

    return f"{document_type}|{document_number}|{dv}"


def parse_number(value: Any) -> float | None:
    value = text(value)

    if value == "":
        return None

    value = value.replace("\u00a0", "").replace(" ", "")

    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")

    elif "," in value:
        value = value.replace(",", ".")

    try:
        number = float(value)
    except ValueError:
        return None

    return number if math.isfinite(number) else None


def parse_birth_year(value: Any) -> int | None:
    if value is None or value == "":
        return None

    if isinstance(value, (int, float)):
        number = float(value)

        if 1900 <= number <= 2100:
            return int(number)

        if 1 <= number <= 100000:
            return (
                datetime(1899, 12, 30)
                + timedelta(days=number)
            ).year

    value = text(value)

    year_match = re.search(
        r"\b(19\d{2}|20\d{2})\b",
        value,
    )

    if year_match:
        return int(year_match.group(1))

    for date_format in (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d-%m-%Y %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
    ):
        try:
            return datetime.strptime(
                value[:19],
                date_format,
            ).year
        except ValueError:
            continue

    number = parse_number(value)

    if number is not None and 1 <= number <= 100000:
        return (
            datetime(1899, 12, 30)
            + timedelta(days=number)
        ).year

    return None


def column_mapping(headers: list[str]) -> dict[str, str | None]:
    aliases = {
        "tipo_documento": ["TIPO_DOCUMENTO"],
        "numero_documento": [
            "NUM_DOCUMENTO",
            "NUMERO_DOCUMENTO",
        ],
        "dv": ["DV"],
        "sexo": ["SEXO"],
        "fecha_nacimiento": [
            "FECHA_NACIMIENTO",
            "ANIO_NACIMIENTO",
            "ANO_NACIMIENTO",
        ],
        "formacion": [
            "NIVEL_FORMACION_ACADEMICO",
            "NIVEL_FORMACION_ACADEMICA",
            "NIVEL_FORMACION_ACADEM",
        ],
        "cargo": ["CARGO_NORMALIZADO"],
        "cargo_texto": [
            "PRINCIPAL_CARGO_ACADEMICO",
        ],
        "programa": [
            "NOMBRE_PRINCIPAL_PROGRAMA",
        ],
        "horas_programa": [
            "TOTAL_HORAS_PRINCIPAL_PROGRAMA",
        ],
        "horas_planta": [
            "NUM_HORAS_PLANTA",
        ],
        "horas_contrata": [
            "NUM_HORAS_CONTRATA",
        ],
        "horas_honorarios": [
            "NUM_HORAS_HONORARIOS",
            "NUM_HORAS_HONORARIO",
        ],
        "especialidad": [
            "NOMBRE_ESPECIALIDAD",
        ],
        "vigencia": ["VIGENCIA"],
        "nacionalidad": ["NACIONALIDAD"],
    }

    return {
        concept: find_column(headers, candidates)
        for concept, candidates in aliases.items()
    }


def build_person_year(
    source: AnnualSource,
) -> list[dict[str, Any]]:
    columns = column_mapping(source.headers)

    required = [
        "tipo_documento",
        "numero_documento",
        "dv",
    ]

    missing = [
        concept
        for concept in required
        if not columns[concept]
    ]

    if missing:
        raise RuntimeError(
            f"{source.year}: faltan columnas {missing}"
        )

    records = []

    for row in source.rows:
        key = canonical_key(row, columns)

        if key is None:
            continue

        birth_year = (
            parse_birth_year(
                row.get(columns["fecha_nacimiento"])
            )
            if columns["fecha_nacimiento"]
            else None
        )

        age = (
            source.year - birth_year
            if birth_year else None
        )

        if age is not None and not 18 <= age <= 100:
            age = None

        hours_plant = (
            parse_number(
                row.get(columns["horas_planta"])
            )
            if columns["horas_planta"]
            else None
        )

        hours_contract = (
            parse_number(
                row.get(columns["horas_contrata"])
            )
            if columns["horas_contrata"]
            else None
        )

        hours_fees = (
            parse_number(
                row.get(columns["horas_honorarios"])
            )
            if columns["horas_honorarios"]
            else None
        )

        components = [
            hours_plant,
            hours_contract,
            hours_fees,
        ]

        total_hours = (
            sum(components)
            if all(
                value is not None
                for value in components
            )
            else None
        )

        program = (
            normalized_business_text(
                row.get(columns["programa"])
            )
            if columns["programa"]
            else ""
        )

        records.append({
            "ANIO": source.year,
            "CLAVE_PERSONA": key,
            "SEXO": normalized_business_text(
                row.get(columns["sexo"])
            ) if columns["sexo"] else "",
            "EDAD": age,
            "PROGRAMA_HOMOLOGADO": program,
            "HORAS_PLANTA": hours_plant,
            "HORAS_CONTRATA": hours_contract,
            "HORAS_HONORARIOS": hours_fees,
            "HORAS_TOTAL": total_hours,
            "FUENTE": source.path.name,
            "SHA256_FUENTE": sha256(source.path),
        })

    return records


def rotation_table(
    annual: dict[int, list[dict[str, Any]]],
) -> list[dict[str, Any]]:
    result = []

    for base_year, next_year in EXPECTED_ROTATION:
        base = {
            row["CLAVE_PERSONA"]
            for row in annual[base_year]
        }

        following = {
            row["CLAVE_PERSONA"]
            for row in annual[next_year]
        }

        remain = base & following
        rotate = base - following
        new = following - base

        rate = len(rotate) / len(base) * 100

        expected = EXPECTED_ROTATION[
            (base_year, next_year)
        ]

        observed = (
            len(base),
            len(following),
            len(remain),
            len(rotate),
            len(new),
        )

        valid = (
            observed == expected[:5]
            and abs(rate - expected[5]) < 0.000001
        )

        result.append({
            "PERIODO": f"{base_year}-{next_year}",
            "DOTACION_BASE": len(base),
            "DOTACION_SIGUIENTE": len(following),
            "PERMANECEN": len(remain),
            "ROTAN": len(rotate),
            "NUEVOS": len(new),
            "TASA_ROTACION": rate,
            "VALIDACION_CONGELADA": (
                "OK" if valid else "REVISAR"
            ),
        })

    return result


def diagnostic() -> dict[str, Any]:
    annual = {}

    for year in EXPECTED_UNIVERSES:
        source = read_source(year)
        records = build_person_year(source)

        keys = {
            row["CLAVE_PERSONA"]
            for row in records
        }

        annual[year] = records

        expected = EXPECTED_UNIVERSES[year]

        if len(keys) != expected:
            raise RuntimeError(
                f"Universo {year}: "
                f"{len(keys)} observado, {expected} esperado."
            )

    rotations = rotation_table(annual)

    if any(
        row["VALIDACION_CONGELADA"] != "OK"
        for row in rotations
    ):
        raise RuntimeError(
            "Los KPI de rotación no coinciden con la línea base."
        )

    all_keys = set().union(
        *(
            {
                row["CLAVE_PERSONA"]
                for row in annual[year]
            }
            for year in annual
        )
    )

    reentries = 0

    for key in all_keys:
        years = [
            year
            for year in sorted(annual)
            if any(
                row["CLAVE_PERSONA"] == key
                for row in annual[year]
            )
        ]

        if any(
            right - left > 1
            for left, right in zip(
                years,
                years[1:],
            )
        ):
            reentries += 1

    if reentries != 13:
        raise RuntimeError(
            f"Reingresos: {reentries} observados, 13 esperados."
        )

    annual_summary = []

    for year, records in annual.items():
        ages = [
            row["EDAD"]
            for row in records
            if row["EDAD"] is not None
        ]

        hours = [
            row["HORAS_TOTAL"]
            for row in records
            if row["HORAS_TOTAL"] is not None
        ]

        sexes = Counter(
            row["SEXO"]
            for row in records
        )

        programs = Counter(
            row["PROGRAMA_HOMOLOGADO"]
            for row in records
        )

        annual_summary.append({
            "ANIO": year,
            "PERSONAS": len(records),
            "EDAD_PROMEDIO": (
                statistics.mean(ages)
                if ages else None
            ),
            "EDAD_MINIMA": min(ages) if ages else None,
            "EDAD_MAXIMA": max(ages) if ages else None,
            "HORAS_TOTALES": (
                sum(hours)
                if hours else None
            ),
            "HORAS_PROMEDIO": (
                statistics.mean(hours)
                if hours else None
            ),
            "HOMBRES": sexes.get("H", 0),
            "MUJERES": sexes.get("M", 0),
            "PROGRAMAS_HOMOLOGADOS": len(programs),
        })

    return {
        "universos": {
            year: len(annual[year])
            for year in annual
        },
        "persona_anio": sum(
            len(records)
            for records in annual.values()
        ),
        "rotacion": rotations,
        "reingresos": reentries,
        "resumen_anual": annual_summary,
    }


def export_excel(output: Path) -> None:
    from openpyxl.cell.cell import MergedCell
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    formation_map = {
        "1": "DOCTORADO",
        "2": "MAGISTER",
        "3": "TITULO PROFESIONAL",
        "4": "LICENCIATURA",
        "5": "TECNICO DE NIVEL SUPERIOR",
        "6": "TITULO TECNICO DE NIVEL MEDIO",
        "7": "LICENCIA MEDIA",
        "8": "SIN INFORMACION",
    }

    cargo_map = {
        "1": "RECTOR / PRO RECTOR",
        "2": "VICERRECTOR DE AREA / SEDE",
        "3": "DECANO / VICE DECANO",
        "4": "ASESOR / CONSEJERO",
        "5": "DIRECTOR O JEFE DE AREA / UNIDAD ACADEMICA",
        "6": "SECRETARIO UNIDAD ACADEMICA",
        "7": "COORDINADOR / ENCARGADO",
        "8": "INVESTIGADOR",
        "9": "DOCENTE",
        "10": "PROFESIONAL / ADMINISTRATIVO",
        "11": "OTRO",
    }

    expected_sex = {
        2022: {"HOMBRE": 60, "MUJER": 11},
        2023: {"HOMBRE": 57, "MUJER": 13},
        2024: {"HOMBRE": 66, "MUJER": 19},
        2025: {"HOMBRE": 85, "MUJER": 39},
        2026: {"HOMBRE": 130, "MUJER": 72},
    }

    expected_hours = {
        2022: 1051.0,
        2023: 1202.0,
        2024: 1703.05,
        2025: 2885.67,
        2026: 3538.58,
    }

    crossed_headers = [
        "ANIO",
        "CLAVE_PERSONA",
        "TIPO_DOCUMENTO",
        "NUM_DOCUMENTO",
        "DV",
        "PRIMER_APELLIDO",
        "SEGUNDO_APELLIDO",
        "NOMBRES",
        "SEXO",
        "SEXO_NORMALIZADO",
        "FECHA_NACIMIENTO",
        "ANIO_NACIMIENTO",
        "EDAD",
        "TRAMO_EDAD",
        "NACIONALIDAD",
        "NIVEL_FORMACION_CODIGO",
        "FORMACION_AGRUPADA",
        "PRINCIPAL_CARGO_ACADEMICO",
        "CARGO_NORMALIZADO",
        "CARGO_AGRUPADO",
        "NOMBRE_PRINCIPAL_PROGRAMA",
        "PROGRAMA_HOMOLOGADO",
        "TOTAL_HORAS_PRINCIPAL_PROGRAMA",
        "NUM_HORAS_PLANTA",
        "NUM_HORAS_CONTRATA",
        "NUM_HORAS_HONORARIOS",
        "TOTAL_HORAS_CONTRACTUALES",
        "TIPO_CONTRACTUAL",
        "TRAMO_HORAS",
        "SUPERA_44_HORAS",
        "NOMBRE_ESPECIALIDAD",
        "VIGENCIA",
        "FUENTE_NORMALIZADA",
        "SHA256_FUENTE",
    ]

    def worksheet_for_year(year: int) -> str:
        return f"{year - 2013:02d}_RAW_{year}"

    def column_map(headers: list[str]) -> dict[str, str | None]:
        columns = column_mapping(headers)
        columns.update({
            "primer_apellido": find_column(headers, ["PRIMER_APELLIDO"]),
            "segundo_apellido": find_column(headers, ["SEGUNDO_APELLIDO"]),
            "nombres": find_column(headers, ["NOMBRES"]),
        })

        return columns

    def get_value(
        row: dict[str, str],
        columns: dict[str, str | None],
        concept: str,
    ) -> str:
        column = columns.get(concept)

        return text(row.get(column)) if column else ""

    def normalized_document_number(value: Any) -> str:
        return re.sub(
            r"\s+",
            "",
            text(value).replace(".", ""),
        ).upper()

    def normalized_dv(value: Any) -> str:
        return re.sub(
            r"\s+",
            "",
            text(value).replace(".", ""),
        ).upper()

    def make_key(
        document_type: Any,
        document_number: Any,
        dv: Any,
    ) -> str:
        parts = [
            text(document_type).upper(),
            normalized_document_number(document_number),
            normalized_dv(dv),
        ]

        if any(part == "" for part in parts):
            return ""

        return "|".join(parts)

    def code(value: Any) -> str:
        value = text(value)

        if value == "":
            return ""

        number = parse_number(value)

        if number is not None and number.is_integer():
            return str(int(number))

        return value.upper()

    def formation(value: Any) -> str:
        return formation_map.get(
            code(value),
            "FORMACION_NO_HOMOLOGABLE",
        )

    def cargo(value: Any) -> str:
        return cargo_map.get(
            code(value),
            "CARGO_NO_HOMOLOGABLE",
        )

    def normalized_sex(value: Any) -> str:
        value = text(value).upper()

        if value == "H":
            return "HOMBRE"

        if value == "M":
            return "MUJER"

        return "SIN_INFORMACION"

    def birth_year(value: Any) -> int | None:
        return parse_birth_year(value)

    def age_group(age: int | None) -> str:
        if age is None:
            return "SIN_DATO"

        if age < 30:
            return "<30"

        if age <= 39:
            return "30-39"

        if age <= 49:
            return "40-49"

        if age <= 59:
            return "50-59"

        return ">=60"

    def numeric(value: Any) -> float | None:
        return parse_number(value)

    def excel_number(value: float | None) -> float | None:
        if value is None:
            return None

        return round(value, 10)

    def contract_type(
        plant: float | None,
        contract: float | None,
        fees: float | None,
    ) -> str:
        components = [plant, contract, fees]

        if any(value is None for value in components):
            return "NO_DETERMINABLE"

        positive = [
            value
            for value in components
            if value and value > 0
        ]

        if not positive:
            return "SIN_HORAS"

        if len(positive) > 1:
            return "COMBINADO"

        if plant and plant > 0:
            return "PLANTA"

        if contract and contract > 0:
            return "CONTRATA"

        return "HONORARIOS"

    def hours_group(total: float | None) -> str:
        if total is None:
            return "SIN_DATO"

        if total < 11:
            return "<11"

        if total <= 22:
            return "11-22"

        if total <= 43:
            return "23-43"

        return "44+"

    def exceeds_44(total: float | None) -> str:
        if total is None:
            return "SIN_DATO"

        return "SI" if total > 44 else "NO"

    def has_formula(sheet: Any) -> int:
        count = 0

        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value

                if cell.data_type == "f":
                    count += 1
                elif isinstance(value, str) and value.startswith("="):
                    count += 1

        return count

    def literal_errors(sheet: Any) -> int:
        errors = {
            "#REF!",
            "#DIV/0!",
            "#VALUE!",
            "#N/A",
            "#NAME?",
            "#NUM!",
            "#NULL!",
        }

        count = 0

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value in errors:
                    count += 1

        return count

    def validation_row(
        control: str,
        expected: Any,
        observed: Any,
        ok: bool,
        detail: str,
    ) -> dict[str, Any]:
        return {
            "CONTROL": control,
            "ESPERADO": expected,
            "OBSERVADO": observed,
            "ESTADO": "OK" if ok else "REVISAR",
            "DETALLE": detail,
        }

    def safe_percent(
        numerator: float | int | None,
        denominator: float | int | None,
    ) -> float | None:
        if numerator is None or denominator in (None, 0):
            return None

        return numerator / denominator * 100

    def safe_ratio(
        numerator: float | int | None,
        denominator: float | int | None,
    ) -> float | None:
        if numerator is None or denominator in (None, 0):
            return None

        return numerator / denominator

    def numeric_values(
        records: list[dict[str, Any]],
        field: str,
    ) -> list[float]:
        return [
            float(record[field])
            for record in records
            if isinstance(record.get(field), (int, float))
            and math.isfinite(float(record[field]))
        ]

    def mean_or_none(values: list[float]) -> float | None:
        return statistics.mean(values) if values else None

    def median_or_none(values: list[float]) -> float | None:
        return statistics.median(values) if values else None

    def min_or_none(values: list[float]) -> float | None:
        return min(values) if values else None

    def max_or_none(values: list[float]) -> float | None:
        return max(values) if values else None

    def rounded(value: Any, digits: int = 6) -> Any:
        if isinstance(value, float):
            return round(value, digits)

        return value

    def period_order(value: Any) -> int:
        match = re.search(r"(20\d{2})", text(value))

        return int(match.group(1)) if match else 0

    def trend_label(
        values: list[float | int | None],
        tolerance: float = 0.000001,
    ) -> str:
        clean = [
            float(value)
            for value in values
            if isinstance(value, (int, float))
            and math.isfinite(float(value))
        ]

        if len(clean) < 2:
            return "SIN_COMPARACION"

        differences = [
            right - left
            for left, right in zip(clean, clean[1:])
        ]

        if all(abs(value) <= tolerance for value in differences):
            return "ESTABLE"

        if all(value > tolerance for value in differences):
            return "CRECIENTE"

        if all(value < -tolerance for value in differences):
            return "DECRECIENTE"

        return "OSCILANTE"

    def movement_label(
        current: float | int | None,
        previous: float | int | None,
        tolerance: float = 0.000001,
    ) -> str:
        if current is None or previous is None:
            return "SIN_COMPARACION"

        difference = float(current) - float(previous)

        if abs(difference) <= tolerance:
            return "ESTABLE"

        return "SUBE" if difference > 0 else "BAJA"

    def rotation_light(variation_pp: float | None) -> str:
        if variation_pp is None:
            return "GRIS"

        if variation_pp < -2.0:
            return "VERDE"

        if variation_pp > 2.0:
            return "ROJO"

        return "AMARILLO"

    def gender_gap_light(gap_pp: float | None) -> str:
        if gap_pp is None:
            return "GRIS"

        if gap_pp <= 10:
            return "VERDE"

        if gap_pp <= 25:
            return "AMARILLO"

        return "ROJO"

    def above_44_light(value: float | None) -> str:
        if value is None:
            return "GRIS"

        if value == 0:
            return "VERDE"

        if value <= 2:
            return "AMARILLO"

        return "ROJO"

    def coverage_light(
        coverage: float | None,
        target: float | None,
    ) -> str:
        if coverage is None or target is None:
            return "GRIS"

        if coverage >= target:
            return "VERDE"

        if coverage >= target - 5:
            return "AMARILLO"

        return "ROJO"

    def hhi_light(value: float | None) -> str:
        if value is None:
            return "GRIS"

        if value < 0.15:
            return "VERDE"

        if value <= 0.25:
            return "AMARILLO"

        return "ROJO"

    def format_value(
        value: Any,
        unit: str,
    ) -> str:
        if value is None:
            return "sin dato"

        if isinstance(value, float):
            number = f"{value:.1f}"
        else:
            number = str(value)

        if unit in {"%", "pp"}:
            return f"{number}{unit}"

        if unit:
            return f"{number} {unit}"

        return number

    def description_text(
        indicator: str,
        period: Any,
        value: Any,
        unit: str,
        previous: Any,
        observation: str,
    ) -> str:
        base = (
            f"En {period}, {indicator} alcanzó "
            f"{format_value(value, unit)}."
        )

        if previous is None:
            comparison = " No existe comparación previa en la serie."
        else:
            movement = movement_label(value, previous).lower()
            gap = (
                float(value) - float(previous)
                if isinstance(value, (int, float))
                and isinstance(previous, (int, float))
                else None
            )
            comparison = (
                f" Respecto de la medición anterior, {movement} "
                f"{format_value(abs(gap), unit)}."
                if gap is not None
                else " Respecto de la medición anterior, no se calculó variación."
            )

        note = f" {observation}" if observation else ""

        return base + comparison + note

    def append_table(
        sheet_name: str,
        headers: list[str],
        rows: list[dict[str, Any]],
    ) -> None:
        sheet = workbook[sheet_name]
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(headers)

        for row in rows:
            sheet.append([
                rounded(row.get(header))
                for header in headers
            ])

    calculation_blocks: dict[str, dict[str, Any]] = {}

    def append_calculation_block(
        sheet: Any,
        title: str,
        headers: list[str],
        rows: list[dict[str, Any]],
    ) -> None:
        if sheet.max_row > 1 or sheet["A1"].value:
            sheet.append([])

        start_row = sheet.max_row + 1
        sheet.append([title])
        header_row = sheet.max_row + 1
        sheet.append(headers)
        data_start = sheet.max_row + 1

        for row in rows:
            sheet.append([
                rounded(row.get(header))
                for header in headers
            ])

        calculation_blocks[title] = {
            "start_row": start_row,
            "header_row": header_row,
            "data_start": data_start,
            "data_end": sheet.max_row,
            "headers": headers,
        }

    def format_numeric_columns(sheet_name: str) -> None:
        sheet = workbook[sheet_name]
        headers = [
            cell.value
            for cell in sheet[1]
        ]

        for column_index, header in enumerate(headers, start=1):
            header_text = text(header).upper()

            if header_text == "":
                continue

            if "HHI" in header_text:
                number_format = "0.000"
            elif any(
                token in header_text
                for token in (
                    "PORC",
                    "TASA",
                    "BRECHA",
                    "VARIACION",
                    "COBERTURA",
                    "PARIDAD",
                    "SEMAFORO",
                )
            ):
                number_format = "0.0"
            elif any(
                token in header_text
                for token in (
                    "HORAS",
                    "EDAD",
                    "VALOR",
                    "META",
                )
            ):
                number_format = "0.0"
            else:
                number_format = "0"

            for row in sheet.iter_rows(
                min_row=2,
                min_col=column_index,
                max_col=column_index,
            ):
                cell = row[0]

                if isinstance(cell.value, (int, float)):
                    cell.number_format = number_format

    def bad_number_count(sheets: list[str]) -> int:
        total = 0

        for sheet_name in sheets:
            sheet = workbook[sheet_name]

            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, float) and not math.isfinite(cell.value):
                        total += 1

        return total

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="1F4E78", bold=True, size=16)
    subtitle_font = Font(color="44546A", bold=True, size=11)
    thin_side = Side(style="thin", color="D9E2F3")
    soft_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )
    card_fill = PatternFill("solid", fgColor="EAF2F8")
    ok_fill = PatternFill("solid", fgColor="D9EAD3")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    danger_fill = PatternFill("solid", fgColor="F4CCCC")
    neutral_fill = PatternFill("solid", fgColor="E7E6E6")
    green_fill = PatternFill("solid", fgColor="C6E0B4")
    yellow_fill = PatternFill("solid", fgColor="FFF2CC")
    red_fill = PatternFill("solid", fgColor="F4CCCC")
    grey_fill = PatternFill("solid", fgColor="E7E6E6")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")

    def table_name(name: str) -> str:
        normalized = re.sub(r"[^A-Za-z0-9_]", "_", name)

        return ("T_" + normalized)[:250]

    def add_excel_table(sheet_name: str, table_prefix: str) -> None:
        sheet = workbook[sheet_name]

        if sheet.max_row < 2 or sheet.max_column < 1:
            return

        reference = (
            f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
        table = Table(
            displayName=table_name(table_prefix),
            ref=reference,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    def autofit(sheet: Any, maximum: int = 40) -> None:
        widths: dict[int, int] = {}

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue

                if cell.value is None:
                    continue

                width = min(
                    len(str(cell.value)) + 2,
                    maximum,
                )
                widths[cell.column] = max(
                    widths.get(cell.column, 0),
                    width,
                )

        for column_index, width in widths.items():
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = max(10, width)

    def apply_common_table_format(
        sheet_name: str,
        maximum_width: int = 40,
        freeze: str = "A2",
        add_table: bool = True,
    ) -> None:
        sheet = workbook[sheet_name]
        sheet.freeze_panes = freeze
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = soft_border

        for cell in sheet[1]:
            if isinstance(cell, MergedCell):
                continue

            cell.font = header_font
            cell.fill = header_fill

        for row_index in range(2, sheet.max_row + 1):
            if row_index % 2 == 0:
                for cell in sheet[row_index]:
                    if not isinstance(cell, MergedCell):
                        cell.fill = PatternFill("solid", fgColor="F8FBFD")

        if add_table:
            add_excel_table(sheet_name, sheet_name)

        autofit(sheet, maximum_width)
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.oddFooter.center.text = "SIIPA v4.0"

    def apply_number_formats(sheet_name: str) -> None:
        sheet = workbook[sheet_name]
        headers = [
            text(cell.value).upper()
            for cell in sheet[1]
        ]

        for column_index, header in enumerate(headers, start=1):
            if "HHI" in header:
                fmt = "0.000"
            elif any(
                token in header
                for token in (
                    "PORC",
                    "TASA",
                    "BRECHA",
                    "VARIACION",
                    "COBERTURA",
                    "PARIDAD",
                )
            ):
                fmt = "0.0"
            elif any(
                token in header
                for token in (
                    "HORAS",
                    "EDAD",
                    "VALOR",
                    "META",
                )
            ):
                fmt = "0.0"
            else:
                fmt = "0"

            for row in sheet.iter_rows(
                min_row=2,
                min_col=column_index,
                max_col=column_index,
            ):
                cell = row[0]

                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt

    def apply_visual_states(sheet_name: str) -> None:
        sheet = workbook[sheet_name]
        headers = [
            text(cell.value).upper()
            for cell in sheet[1]
        ]
        fills = {
            "VERDE": green_fill,
            "AMARILLO": yellow_fill,
            "ROJO": red_fill,
            "GRIS": grey_fill,
            "OK": green_fill,
            "REVISAR": yellow_fill,
            "BLOQUEANTE": red_fill,
            "CUMPLE": green_fill,
            "DESCRIPTIVO": grey_fill,
        }

        for column_index, header in enumerate(headers, start=1):
            if not any(
                token in header
                for token in (
                    "SEMAFORO",
                    "ESTADO",
                    "BRECHA",
                    "TENDENCIA",
                )
            ):
                continue

            for row in sheet.iter_rows(
                min_row=2,
                min_col=column_index,
                max_col=column_index,
            ):
                cell = row[0]
                value = text(cell.value).upper()

                if value in fills:
                    cell.fill = fills[value]
                elif isinstance(cell.value, (int, float)) and "BRECHA" in header:
                    magnitude = abs(float(cell.value))

                    if magnitude <= 5:
                        cell.fill = green_fill
                    elif magnitude <= 15:
                        cell.fill = yellow_fill
                    elif magnitude <= 30:
                        cell.fill = orange_fill
                    else:
                        cell.fill = red_fill

    def configure_sheet_print(sheet_name: str) -> None:
        sheet = workbook[sheet_name]
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.oddHeader.center.text = "SIIPA v4.0"

    def build_cover() -> None:
        sheet = workbook["00_PORTADA"]
        sheet.delete_rows(1, sheet.max_row)
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:H1")
        sheet["A1"] = "SISTEMA INTEGRADO DE INDICADORES DE PERSONAL ACADÉMICO"
        sheet["A1"].font = Font(color="1F4E78", bold=True, size=20)
        sheet["A1"].alignment = Alignment(horizontal="center")
        sheet.merge_cells("A2:H2")
        sheet["A2"] = "SIIPA — Línea Base Histórica 2022-2026"
        sheet["A2"].font = Font(color="44546A", bold=True, size=14)
        sheet["A2"].alignment = Alignment(horizontal="center")

        cover_rows = [
            ("Versión", "4.0"),
            ("Fecha de generación", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("Fuente", "Personal Académico SIES"),
            ("Universos", "2022: 71; 2023: 70; 2024: 85; 2025: 124; 2026: 202"),
            ("Persona-año", 552),
            ("Periodos de rotación", 4),
            ("Advertencia", "Análisis descriptivo censal."),
            ("Nota rotación", "La rotación no implica necesariamente desvinculación definitiva."),
            ("Restricción", "Los datos individuales son de uso restringido."),
        ]

        start = 5

        for index, (label, value) in enumerate(cover_rows, start=start):
            sheet.cell(index, 1, label)
            sheet.cell(index, 2, value)

        sheet.cell(16, 1, "KPI congelados de rotación")
        sheet.cell(16, 1).font = subtitle_font
        row_index = 17

        for row in stability_rows:
            sheet.cell(row_index, 1, row["PERIODO"])
            sheet.cell(row_index, 2, row["TASA_ROTACION"] / 100)
            sheet.cell(row_index, 2).number_format = "0.0%"
            row_index += 1

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = soft_border

        for cell in sheet["A"]:
            if not isinstance(cell, MergedCell) and cell.row >= start:
                cell.font = Font(bold=True, color="1F4E78")

        autofit(sheet, 50)
        configure_sheet_print("00_PORTADA")

    def write_card(
        sheet: Any,
        row: int,
        col: int,
        title: str,
        value: Any,
        change: Any,
        light: str,
        description: str,
    ) -> None:
        labels = [title, value, change, light, description]

        for offset, item in enumerate(labels):
            cell = sheet.cell(row + offset, col, item)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = soft_border

            if offset == 0:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
            elif offset == 3:
                cell.fill = {
                    "VERDE": green_fill,
                    "AMARILLO": yellow_fill,
                    "ROJO": red_fill,
                    "GRIS": grey_fill,
                }.get(text(item).upper(), card_fill)
            else:
                cell.fill = card_fill

        sheet.row_dimensions[row + 4].height = 42

    def find_kpi(
        indicator: str,
        period: Any,
    ) -> dict[str, Any] | None:
        for row in kpi_rows:
            if row["INDICADOR"] == indicator and row["ANIO_PERIODO"] == period:
                return row

        return None

    def build_dashboard() -> None:
        sheet = workbook["01_RESUMEN_EJECUTIVO"]
        sheet.delete_rows(1, sheet.max_row)
        sheet.sheet_view.showGridLines = False
        sheet["A1"] = "Resumen ejecutivo SIIPA v4.0"
        sheet["A1"].font = title_font
        sheet["A2"] = "Indicadores descriptivos de Personal Académico SIES 2022-2026"
        sheet["A2"].font = subtitle_font

        cards = [
            ("Dotación 2026", find_kpi("DOTACION_TOTAL", 2026), "personas"),
            ("Variación dotación 2025-2026", find_kpi("DOTACION_TOTAL", 2026), "personas"),
            ("Rotación 2025-2026", find_kpi("TASA_ROTACION", "2025-2026"), "%"),
            ("Permanencia 2025-2026", find_kpi("TASA_PERMANENCIA", "2025-2026"), "%"),
            ("Mujeres 2026", find_kpi("PORC_MUJERES", 2026), "%"),
            ("Edad promedio 2026", find_kpi("EDAD_PROMEDIO", 2026), "años"),
            ("Horas totales 2026", find_kpi("HORAS_TOTALES", 2026), "horas"),
            ("Horas promedio 2026", find_kpi("HORAS_PROMEDIO", 2026), "horas"),
            ("Magíster 2026", find_kpi("PORC_MAGISTER", 2026), "%"),
            ("Docentes 2026", find_kpi("PORC_DOCENTES", 2026), "%"),
            ("Reingresos históricos", {"VALOR": 13, "BRECHA_ABSOLUTA": None, "SEMAFORO": "GRIS", "DESCRIPCION_DINAMICA": "Entre 2022 y 2026 se observan 13 reingresos agregados."}, "personas"),
            ("Programas 2026", find_kpi("NUM_PROGRAMAS", 2026), "programas"),
        ]
        card_positions = [
            (4, 1), (4, 3), (4, 5), (4, 7),
            (10, 1), (10, 3), (10, 5), (10, 7),
            (16, 1), (16, 3), (16, 5), (16, 7),
        ]

        for (title, kpi, unit), (row, col) in zip(cards, card_positions):
            if kpi is None:
                value = "SIN_DATO"
                change = "SIN_DATO"
                light = "GRIS"
                description = "Indicador no disponible."
            else:
                value = format_value(kpi.get("VALOR"), unit)
                change = (
                    format_value(kpi.get("BRECHA_ABSOLUTA"), unit)
                    if kpi.get("BRECHA_ABSOLUTA") is not None
                    else "sin comparación"
                )
                light = text(kpi.get("SEMAFORO") or "GRIS")
                description = text(kpi.get("DESCRIPCION_DINAMICA"))

            write_card(sheet, row, col, title, value, change, light, description)

        red_alerts = [
            row
            for row in kpi_rows
            if row.get("SEMAFORO") == "ROJO"
        ][:6]
        sheet["A23"] = "Alertas y hallazgos"
        sheet["A23"].font = subtitle_font
        alert_row = 24

        if not red_alerts:
            sheet.cell(alert_row, 1, "No se registran semáforos rojos en los KPI agregados.")
        else:
            for alert in red_alerts:
                sheet.cell(
                    alert_row,
                    1,
                    f"{alert['ANIO_PERIODO']} · {alert['INDICADOR']}: "
                    f"{format_value(alert['VALOR'], alert['UNIDAD'])}. "
                    f"{alert['OBSERVACION']}",
                )
                alert_row += 1

        dot_2025 = find_kpi("DOTACION_TOTAL", 2025)
        dot_2026 = find_kpi("DOTACION_TOTAL", 2026)
        rot_2025_2026 = find_kpi("TASA_ROTACION", "2025-2026")
        women_2026 = find_kpi("PORC_MUJERES", 2026)
        age_2026 = find_kpi("EDAD_PROMEDIO", 2026)
        hours_2026 = find_kpi("HORAS_TOTALES", 2026)
        magister_2026 = find_kpi("PORC_MAGISTER", 2026)
        sheet["A31"] = "Síntesis dinámica"
        sheet["A31"].font = subtitle_font
        sheet["A32"] = (
            f"Entre 2025 y 2026 la dotación pasó de "
            f"{dot_2025['VALOR']} a {dot_2026['VALOR']} personas. "
            f"La rotación 2025-2026 fue {format_value(rot_2025_2026['VALOR'], '%')}. "
            f"En 2026, las mujeres representaron {format_value(women_2026['VALOR'], '%')}; "
            f"la edad promedio fue {format_value(age_2026['VALOR'], 'años')} y "
            f"las horas contractuales sumaron {format_value(hours_2026['VALOR'], 'horas')}. "
            f"El porcentaje con magíster fue {format_value(magister_2026['VALOR'], '%')}. "
            "El análisis es descriptivo censal y no atribuye causalidad."
        )
        sheet["A32"].alignment = Alignment(wrap_text=True, vertical="top")

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                cell.border = soft_border

        sheet.row_dimensions[32].height = 56
        autofit(sheet, 34)
        configure_sheet_print("01_RESUMEN_EJECUTIVO")

    def append_chart_block(title: str, headers: list[str], rows: list[dict[str, Any]]) -> None:
        append_calculation_block(workbook["92_CALCULOS"], title, headers, rows)

    def chart_block_ref(
        block_name: str,
        columns: list[int],
    ) -> tuple[Reference, Reference]:
        block = calculation_blocks[block_name]
        min_col = min(columns)
        max_col = max(columns)
        data = Reference(
            workbook["92_CALCULOS"],
            min_col=min_col,
            max_col=max_col,
            min_row=block["header_row"],
            max_row=block["data_end"],
        )
        categories = Reference(
            workbook["92_CALCULOS"],
            min_col=1,
            max_col=1,
            min_row=block["data_start"],
            max_row=block["data_end"],
        )

        return data, categories

    def add_chart(
        name: str,
        chart_type: str,
        block_name: str,
        columns: list[int],
        anchor: str,
    ) -> dict[str, Any]:
        chart = LineChart() if chart_type == "line" else BarChart()

        if chart_type == "bar":
            chart.type = "bar"
        elif chart_type == "stacked":
            chart.type = "col"
            chart.grouping = "stacked"
            chart.overlap = 100
        else:
            chart.type = "col"

        data, categories = chart_block_ref(block_name, columns)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.title = name
        chart.height = 7
        chart.width = 12
        chart.legend.position = "b"
        chart.dataLabels = DataLabelList()
        workbook["01_RESUMEN_EJECUTIVO"].add_chart(chart, anchor)
        block = calculation_blocks[block_name]

        return {
            "NOMBRE": name,
            "TIPO": chart_type,
            "HOJA_DESTINO": "01_RESUMEN_EJECUTIVO",
            "CELDA_ANCLA": anchor,
            "RANGO_CATEGORIAS": (
                f"92_CALCULOS!A{block['data_start']}:"
                f"A{block['data_end']}"
            ),
            "RANGO_SERIES": (
                f"92_CALCULOS!{get_column_letter(min(columns))}{block['header_row']}:"
                f"{get_column_letter(max(columns))}{block['data_end']}"
            ),
            "SERIES": len(columns),
            "ESTADO": "OK",
        }

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    sources = {
        year: read_source(year)
        for year in EXPECTED_UNIVERSES
    }

    source_hashes = {
        year: sha256(source.path)
        for year, source in sources.items()
    }

    source_columns = {
        year: column_map(source.headers)
        for year, source in sources.items()
    }

    workbook = Workbook()

    first_sheet = workbook.active
    first_sheet.title = SIIPA_SHEET_NAMES[0]
    first_sheet["A1"] = SIIPA_SHEET_NAMES[0]

    for sheet_name in SIIPA_SHEET_NAMES[1:]:
        sheet = workbook.create_sheet(sheet_name)
        sheet["A1"] = sheet_name

    for year, source in sources.items():
        sheet = workbook[worksheet_for_year(year)]
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(source.headers)

        for row in source.rows:
            sheet.append([
                row.get(header, "")
                for header in source.headers
            ])

    crossed_records = []

    for year, source in sources.items():
        columns = source_columns[year]

        for row in source.rows:
            document_type = get_value(row, columns, "tipo_documento")
            document_number = get_value(row, columns, "numero_documento")
            dv = get_value(row, columns, "dv")
            year_of_birth = birth_year(
                get_value(row, columns, "fecha_nacimiento")
            )
            age = (
                year - year_of_birth
                if year_of_birth is not None
                else None
            )
            hours_program = numeric(
                get_value(row, columns, "horas_programa")
            )
            hours_plant = numeric(
                get_value(row, columns, "horas_planta")
            )
            hours_contract = numeric(
                get_value(row, columns, "horas_contrata")
            )
            hours_fees = numeric(
                get_value(row, columns, "horas_honorarios")
            )
            components = [
                hours_plant,
                hours_contract,
                hours_fees,
            ]
            total_hours = (
                sum(components)
                if all(value is not None for value in components)
                else None
            )
            program = get_value(row, columns, "programa")

            crossed_records.append({
                "ANIO": year,
                "CLAVE_PERSONA": make_key(
                    document_type,
                    document_number,
                    dv,
                ),
                "TIPO_DOCUMENTO": document_type,
                "NUM_DOCUMENTO": document_number,
                "DV": dv,
                "PRIMER_APELLIDO": get_value(
                    row,
                    columns,
                    "primer_apellido",
                ),
                "SEGUNDO_APELLIDO": get_value(
                    row,
                    columns,
                    "segundo_apellido",
                ),
                "NOMBRES": get_value(row, columns, "nombres"),
                "SEXO": get_value(row, columns, "sexo"),
                "SEXO_NORMALIZADO": normalized_sex(
                    get_value(row, columns, "sexo")
                ),
                "FECHA_NACIMIENTO": get_value(
                    row,
                    columns,
                    "fecha_nacimiento",
                ),
                "ANIO_NACIMIENTO": year_of_birth,
                "EDAD": age,
                "TRAMO_EDAD": age_group(age),
                "NACIONALIDAD": get_value(
                    row,
                    columns,
                    "nacionalidad",
                ),
                "NIVEL_FORMACION_CODIGO": get_value(
                    row,
                    columns,
                    "formacion",
                ),
                "FORMACION_AGRUPADA": formation(
                    get_value(row, columns, "formacion")
                ),
                "PRINCIPAL_CARGO_ACADEMICO": get_value(
                    row,
                    columns,
                    "cargo_texto",
                ),
                "CARGO_NORMALIZADO": get_value(
                    row,
                    columns,
                    "cargo",
                ),
                "CARGO_AGRUPADO": cargo(
                    get_value(row, columns, "cargo")
                ),
                "NOMBRE_PRINCIPAL_PROGRAMA": program,
                "PROGRAMA_HOMOLOGADO": normalized_business_text(program),
                "TOTAL_HORAS_PRINCIPAL_PROGRAMA": excel_number(
                    hours_program
                ),
                "NUM_HORAS_PLANTA": excel_number(hours_plant),
                "NUM_HORAS_CONTRATA": excel_number(hours_contract),
                "NUM_HORAS_HONORARIOS": excel_number(hours_fees),
                "TOTAL_HORAS_CONTRACTUALES": excel_number(total_hours),
                "TIPO_CONTRACTUAL": contract_type(
                    hours_plant,
                    hours_contract,
                    hours_fees,
                ),
                "TRAMO_HORAS": hours_group(total_hours),
                "SUPERA_44_HORAS": exceeds_44(total_hours),
                "NOMBRE_ESPECIALIDAD": get_value(
                    row,
                    columns,
                    "especialidad",
                ),
                "VIGENCIA": get_value(row, columns, "vigencia"),
                "FUENTE_NORMALIZADA": source.path.name,
                "SHA256_FUENTE": source_hashes[year],
            })

    crossed_sheet = workbook["08_DATOS_CRUZADOS"]
    crossed_sheet.delete_rows(1, crossed_sheet.max_row)
    crossed_sheet.append(crossed_headers)

    for record in crossed_records:
        crossed_sheet.append([
            record[header]
            for header in crossed_headers
        ])

    total_person_year = len(crossed_records)
    incomplete_keys = sum(
        1
        for record in crossed_records
        if not record["CLAVE_PERSONA"]
    )
    duplicate_person_year = (
        total_person_year
        - len({
            (
                record["ANIO"],
                record["CLAVE_PERSONA"],
            )
            for record in crossed_records
        })
    )
    formulas_in_data = (
        has_formula(crossed_sheet)
        + sum(
            has_formula(workbook[worksheet_for_year(year)])
            for year in sources
        )
    )
    used_paths = [
        str(source.path)
        for source in sources.values()
    ] + [str(output)]

    annual_records = {
        year: [
            record
            for record in crossed_records
            if record["ANIO"] == year
        ]
        for year in EXPECTED_UNIVERSES
    }
    annual_keys = {
        year: {
            record["CLAVE_PERSONA"]
            for record in records
        }
        for year, records in annual_records.items()
    }
    reentries_by_year = {}

    for year in EXPECTED_UNIVERSES:
        previous_year = year - 1
        previous_keys = annual_keys.get(previous_year, set())
        earlier_keys = set().union(
            *[
                keys
                for earlier, keys in annual_keys.items()
                if earlier < previous_year
            ]
        ) if any(
            earlier < previous_year
            for earlier in annual_keys
        ) else set()

        reentries_by_year[year] = len(
            annual_keys[year]
            - previous_keys
            & earlier_keys
        )

    annual_summary = []
    diversity_rows = []
    hours_rows = []
    category_rows = []
    quality_rows = []
    program_summary = []
    formation_summary = []
    cargo_summary = []
    contract_summary = []
    age_summary = []
    gender_summary = []

    previous_annual = None
    previous_category = {}

    quality_targets = {
        "clave documental completa": 100.0,
        "duplicados persona-año": 100.0,
        "sexo interpretable": 95.0,
        "fecha de nacimiento presente": 100.0,
        "edad calculable": 100.0,
        "edad plausible": 100.0,
        "horas planta calculables": 100.0,
        "horas contrata calculables": 100.0,
        "horas honorarios calculables": 100.0,
        "horas totales calculables": 100.0,
        "formación homologable": 95.0,
        "cargo homologable": 95.0,
        "programa informado": 95.0,
        "especialidad informada": None,
        "vigencia informada": 95.0,
        "documento sin espacios": 100.0,
        "DV interpretable": 100.0,
    }

    for year, records in annual_records.items():
        total = len(records)
        sexes = Counter(
            record["SEXO_NORMALIZADO"]
            for record in records
        )
        ages = numeric_values(records, "EDAD")
        hours = numeric_values(records, "TOTAL_HORAS_CONTRACTUALES")
        plant_hours = sum(
            numeric_values(records, "NUM_HORAS_PLANTA")
        )
        contract_hours = sum(
            numeric_values(records, "NUM_HORAS_CONTRATA")
        )
        fees_hours = sum(
            numeric_values(records, "NUM_HORAS_HONORARIOS")
        )
        age_groups = Counter(
            record["TRAMO_EDAD"]
            for record in records
        )
        hour_groups = Counter(
            record["TRAMO_HORAS"]
            for record in records
        )
        formations = Counter(
            record["FORMACION_AGRUPADA"]
            for record in records
        )
        cargos = Counter(
            record["CARGO_AGRUPADO"]
            for record in records
        )
        contracts = Counter(
            record["TIPO_CONTRACTUAL"]
            for record in records
        )
        programs = Counter(
            record["PROGRAMA_HOMOLOGADO"] or "SIN_INFORMACION"
            for record in records
        )
        women_pct = safe_percent(sexes.get("MUJER", 0), total)
        men_pct = safe_percent(sexes.get("HOMBRE", 0), total)
        gender_gap = (
            abs(men_pct - women_pct)
            if men_pct is not None and women_pct is not None
            else None
        )
        parity = (
            safe_percent(
                min(sexes.get("HOMBRE", 0), sexes.get("MUJER", 0)),
                max(sexes.get("HOMBRE", 0), sexes.get("MUJER", 0)),
            )
            if max(sexes.get("HOMBRE", 0), sexes.get("MUJER", 0)) > 0
            else None
        )
        age_average = mean_or_none(ages)
        total_hours = sum(hours)
        hours_average = mean_or_none(hours)
        hours_over_44 = sum(
            1
            for record in records
            if record["TOTAL_HORAS_CONTRACTUALES"] is not None
            and record["TOTAL_HORAS_CONTRACTUALES"] > 44
        )
        hours_over_44_pct = safe_percent(hours_over_44, total)
        top_program_counts = [
            count
            for _, count in programs.most_common()
        ]
        top_1 = safe_percent(top_program_counts[0], total) if top_program_counts else None
        top_3 = safe_percent(sum(top_program_counts[:3]), total) if top_program_counts else None
        hhi = sum(
            (count / total) ** 2
            for count in top_program_counts
        ) if total else None
        directivos = sum(
            count
            for category, count in cargos.items()
            if category in {
                "RECTOR / PRO RECTOR",
                "VICERRECTOR DE AREA / SEDE",
                "DECANO / VICE DECANO",
                "DIRECTOR O JEFE DE AREA / UNIDAD ACADEMICA",
            }
        )

        annual = {
            "ANIO": year,
            "DOTACION_TOTAL": total,
            "HOMBRES": sexes.get("HOMBRE", 0),
            "MUJERES": sexes.get("MUJER", 0),
            "SIN_INFORMACION_SEXO": sexes.get("SIN_INFORMACION", 0),
            "PORC_HOMBRES": men_pct,
            "PORC_MUJERES": women_pct,
            "BRECHA_GENERO_PP": gender_gap,
            "INDICE_PARIDAD_GENERO": parity,
            "EDAD_PROMEDIO": age_average,
            "EDAD_MEDIANA": median_or_none(ages),
            "EDAD_MINIMA": min_or_none(ages),
            "EDAD_MAXIMA": max_or_none(ages),
            "PORC_MENORES_30": safe_percent(age_groups.get("<30", 0), total),
            "PORC_30_39": safe_percent(age_groups.get("30-39", 0), total),
            "PORC_40_49": safe_percent(age_groups.get("40-49", 0), total),
            "PORC_50_59": safe_percent(age_groups.get("50-59", 0), total),
            "PORC_60_MAS": safe_percent(age_groups.get(">=60", 0), total),
            "HORAS_TOTALES": total_hours,
            "HORAS_PLANTA": plant_hours,
            "HORAS_CONTRATA": contract_hours,
            "HORAS_HONORARIOS": fees_hours,
            "HORAS_PROMEDIO": hours_average,
            "HORAS_MEDIANA": median_or_none(hours),
            "HORAS_MINIMA": min_or_none(hours),
            "HORAS_MAXIMA": max_or_none(hours),
            "PERSONAS_PLANTA": contracts.get("PLANTA", 0),
            "PERSONAS_CONTRATA": contracts.get("CONTRATA", 0),
            "PERSONAS_HONORARIOS": contracts.get("HONORARIOS", 0),
            "PERSONAS_COMBINADAS": contracts.get("COMBINADO", 0),
            "PERSONAS_SIN_HORAS": contracts.get("SIN_HORAS", 0),
            "PERSONAS_NO_DETERMINABLES": contracts.get("NO_DETERMINABLE", 0),
            "PERSONAS_MAS_44_HORAS": hours_over_44,
            "PORC_MAS_44_HORAS": hours_over_44_pct,
            "PORC_HORAS_PLANTA": safe_percent(plant_hours, total_hours),
            "PORC_HORAS_HONORARIOS": safe_percent(fees_hours, total_hours),
            "NUM_PROGRAMAS": len(programs),
            "CONCENTRACION_TOP_1": top_1,
            "CONCENTRACION_TOP_3": top_3,
            "INDICE_HHI_PROGRAMAS": hhi,
            "PORC_DOCTORADO": safe_percent(formations.get("DOCTORADO", 0), total),
            "PORC_MAGISTER": safe_percent(formations.get("MAGISTER", 0), total),
            "PORC_TITULO_PROFESIONAL": safe_percent(
                formations.get("TITULO PROFESIONAL", 0),
                total,
            ),
            "PORC_TECNICO_NIVEL_SUPERIOR": safe_percent(
                formations.get("TECNICO DE NIVEL SUPERIOR", 0),
                total,
            ),
            "PORC_FORMACION_POSTGRADO": safe_percent(
                formations.get("DOCTORADO", 0)
                + formations.get("MAGISTER", 0),
                total,
            ),
            "PORC_DOCENTES": safe_percent(cargos.get("DOCENTE", 0), total),
            "PORC_DIRECTIVOS": safe_percent(directivos, total),
            "PORC_COORDINADORES": safe_percent(
                cargos.get("COORDINADOR / ENCARGADO", 0),
                total,
            ),
            "REINGRESOS": reentries_by_year[year],
        }

        annual_summary.append(annual)

        diversity = {
            key: annual.get(key)
            for key in (
                "ANIO",
                "DOTACION_TOTAL",
                "HOMBRES",
                "MUJERES",
                "SIN_INFORMACION_SEXO",
                "PORC_HOMBRES",
                "PORC_MUJERES",
                "BRECHA_GENERO_PP",
                "INDICE_PARIDAD_GENERO",
                "EDAD_PROMEDIO",
                "EDAD_MEDIANA",
                "EDAD_MINIMA",
                "EDAD_MAXIMA",
                "PORC_MENORES_30",
                "PORC_30_39",
                "PORC_40_49",
                "PORC_50_59",
                "PORC_60_MAS",
            )
        }
        diversity["VARIACION_PORC_MUJERES_PP"] = (
            annual["PORC_MUJERES"] - previous_annual["PORC_MUJERES"]
            if previous_annual else None
        )
        diversity["VARIACION_EDAD_PROMEDIO"] = (
            annual["EDAD_PROMEDIO"] - previous_annual["EDAD_PROMEDIO"]
            if previous_annual else None
        )
        diversity["SEMAFORO_BRECHA_GENERO"] = gender_gap_light(gender_gap)
        diversity["DESCRIPCION"] = (
            f"En {year}, las mujeres representan "
            f"{format_value(annual['PORC_MUJERES'], '%')} y la brecha "
            f"agregada de representación es {format_value(gender_gap, 'pp')}. "
            "El indicador es descriptivo y no establece causalidad."
        )
        diversity_rows.append(diversity)
        gender_summary.append({
            "ANIO": year,
            "HOMBRES": annual["HOMBRES"],
            "MUJERES": annual["MUJERES"],
            "SIN_INFORMACION_SEXO": annual["SIN_INFORMACION_SEXO"],
            "PORC_HOMBRES": annual["PORC_HOMBRES"],
            "PORC_MUJERES": annual["PORC_MUJERES"],
        })
        age_summary.append({
            "ANIO": year,
            "EDAD_PROMEDIO": annual["EDAD_PROMEDIO"],
            "EDAD_MEDIANA": annual["EDAD_MEDIANA"],
            "PORC_MENORES_30": annual["PORC_MENORES_30"],
            "PORC_60_MAS": annual["PORC_60_MAS"],
        })

        hours_row = {
            key: annual.get(key)
            for key in (
                "ANIO",
                "DOTACION_TOTAL",
                "HORAS_TOTALES",
                "HORAS_PLANTA",
                "HORAS_CONTRATA",
                "HORAS_HONORARIOS",
                "HORAS_PROMEDIO",
                "HORAS_MEDIANA",
                "HORAS_MINIMA",
                "HORAS_MAXIMA",
                "PERSONAS_PLANTA",
                "PERSONAS_CONTRATA",
                "PERSONAS_HONORARIOS",
                "PERSONAS_COMBINADAS",
                "PERSONAS_SIN_HORAS",
                "PERSONAS_NO_DETERMINABLES",
                "PERSONAS_MAS_44_HORAS",
                "PORC_MAS_44_HORAS",
                "PORC_HORAS_PLANTA",
                "PORC_HORAS_HONORARIOS",
            )
        }
        hours_row["PERSONAS"] = hours_row.pop("DOTACION_TOTAL")
        hours_row["VARIACION_HORAS_TOTALES"] = (
            annual["HORAS_TOTALES"] - previous_annual["HORAS_TOTALES"]
            if previous_annual else None
        )
        hours_row["VARIACION_HORAS_PROMEDIO"] = (
            annual["HORAS_PROMEDIO"] - previous_annual["HORAS_PROMEDIO"]
            if previous_annual else None
        )
        hours_row["SEMAFORO_MAS_44"] = above_44_light(hours_over_44_pct)
        hours_row["DESCRIPCION"] = (
            f"En {year}, las horas contractuales totales fueron "
            f"{format_value(total_hours, 'horas')}. "
            f"{format_value(hours_over_44_pct, '%')} de las personas "
            "supera 44 horas contractuales; este control no asume "
            "mejor resultado institucional."
        )
        hours_rows.append(hours_row)

        for dimension, counter in (
            ("FORMACION", formations),
            ("CARGO", cargos),
            ("TIPO_CONTRACTUAL", contracts),
            ("TRAMO_EDAD", age_groups),
            ("TRAMO_HORAS", hour_groups),
            ("PROGRAMA", programs),
        ):
            for category, count in sorted(counter.items()):
                percentage = safe_percent(count, total)
                previous = previous_category.get((dimension, category))
                variation_count = (
                    count - previous["PERSONAS"]
                    if previous else None
                )
                variation_pct = (
                    percentage - previous["PORCENTAJE"]
                    if previous else None
                )
                category_row = {
                    "ANIO": year,
                    "DIMENSION": dimension,
                    "CATEGORIA": category,
                    "PERSONAS": count,
                    "PORCENTAJE": percentage,
                    "VARIACION_PERSONAS": variation_count,
                    "VARIACION_PORCENTAJE_PP": variation_pct,
                    "TENDENCIA": movement_label(count, previous["PERSONAS"]) if previous else "SIN_COMPARACION",
                    "SEMAFORO": "GRIS",
                    "DESCRIPCION": (
                        f"En {year}, la categoría {category} de {dimension} "
                        f"concentra {count} personas "
                        f"({format_value(percentage, '%')}). "
                        "La lectura es descriptiva y no causal."
                    ),
                }
                category_rows.append(category_row)
                previous_category[(dimension, category)] = category_row

                if dimension == "FORMACION":
                    formation_summary.append(category_row)
                elif dimension == "CARGO":
                    cargo_summary.append(category_row)
                elif dimension == "TIPO_CONTRACTUAL":
                    contract_summary.append(category_row)
                elif dimension == "PROGRAMA":
                    program_summary.append(category_row)

        duplicate_keys_year = total - len(annual_keys[year])
        quality_definitions = [
            (
                "clave documental completa",
                sum(1 for record in records if record["CLAVE_PERSONA"]),
                "Identidad documental completa.",
            ),
            (
                "duplicados persona-año",
                total - duplicate_keys_year,
                "Ausencia de claves duplicadas por año.",
            ),
            (
                "sexo interpretable",
                sum(
                    1
                    for record in records
                    if record["SEXO_NORMALIZADO"] in {"HOMBRE", "MUJER"}
                ),
                "Sexo homologado a HOMBRE o MUJER.",
            ),
            (
                "fecha de nacimiento presente",
                sum(1 for record in records if record["FECHA_NACIMIENTO"]),
                "Fecha de nacimiento informada.",
            ),
            (
                "edad calculable",
                len(ages),
                "Edad calculada al 31 de diciembre.",
            ),
            (
                "edad plausible",
                sum(1 for age in ages if 18 <= age <= 100),
                "Edad entre 18 y 100 años.",
            ),
            (
                "horas planta calculables",
                len(numeric_values(records, "NUM_HORAS_PLANTA")),
                "Horas planta numéricas.",
            ),
            (
                "horas contrata calculables",
                len(numeric_values(records, "NUM_HORAS_CONTRATA")),
                "Horas contrata numéricas.",
            ),
            (
                "horas honorarios calculables",
                len(numeric_values(records, "NUM_HORAS_HONORARIOS")),
                "Horas honorarios numéricas.",
            ),
            (
                "horas totales calculables",
                len(hours),
                "Total contractual calculado desde componentes.",
            ),
            (
                "formación homologable",
                sum(
                    1
                    for record in records
                    if record["FORMACION_AGRUPADA"] != "FORMACION_NO_HOMOLOGABLE"
                ),
                "Código de formación reconocido.",
            ),
            (
                "cargo homologable",
                sum(
                    1
                    for record in records
                    if record["CARGO_AGRUPADO"] != "CARGO_NO_HOMOLOGABLE"
                ),
                "Código de cargo reconocido.",
            ),
            (
                "programa informado",
                sum(1 for record in records if record["PROGRAMA_HOMOLOGADO"]),
                "Programa principal informado.",
            ),
            (
                "especialidad informada",
                sum(1 for record in records if record["NOMBRE_ESPECIALIDAD"]),
                "Especialidad informada; control descriptivo.",
            ),
            (
                "vigencia informada",
                sum(1 for record in records if record["VIGENCIA"]),
                "Vigencia informada.",
            ),
            (
                "documento sin espacios",
                sum(
                    1
                    for record in records
                    if not re.search(r"\s", text(record["NUM_DOCUMENTO"]))
                ),
                "Número documental sin espacios internos.",
            ),
            (
                "DV interpretable",
                sum(
                    1
                    for record in records
                    if re.fullmatch(r"[0-9Kk]+", text(record["DV"]))
                ),
                "DV alfanumérico interpretable.",
            ),
        ]

        for control, conformes, source_detail in quality_definitions:
            target = quality_targets[control]
            coverage = safe_percent(conformes, total)
            breach = (
                coverage - target
                if coverage is not None and target is not None
                else None
            )
            quality_rows.append({
                "ANIO": year,
                "CONTROL": control,
                "UNIVERSO": total,
                "CONFORMES": conformes,
                "NO_CONFORMES": total - conformes,
                "COBERTURA_PCT": coverage,
                "BRECHA_META_PP": breach,
                "ESTADO": "DESCRIPTIVO" if target is None else (
                    "CUMPLE" if coverage is not None and coverage >= target else "REVISAR"
                ),
                "SEMAFORO": "GRIS" if target is None else coverage_light(coverage, target),
                "DESCRIPCION": (
                    f"En {year}, el control '{control}' presenta "
                    f"{format_value(coverage, '%')} de cobertura."
                ),
                "FUENTE_VALIDACION": source_detail,
            })

        previous_annual = annual

    stability_rows = []
    previous_rate = None

    for base_year, next_year in EXPECTED_ROTATION:
        expected = EXPECTED_ROTATION[(base_year, next_year)]
        dot_base, dot_next, remain, rotate, new, rate = expected
        permanence_rate = remain / dot_base * 100
        variation = (
            rate - previous_rate
            if previous_rate is not None
            else None
        )
        trend = movement_label(rate, previous_rate)

        stability_rows.append({
            "PERIODO": f"{base_year}-{next_year}",
            "DOTACION_BASE": dot_base,
            "DOTACION_SIGUIENTE": dot_next,
            "PERMANECEN": remain,
            "ROTAN": rotate,
            "NUEVOS": new,
            "TASA_ROTACION": rate,
            "TASA_PERMANENCIA": permanence_rate,
            "VARIACION_ROTACION_PP": variation,
            "TENDENCIA_ROTACION": trend,
            "SEMAFORO_ROTACION": rotation_light(variation),
            "REINGRESOS_RELACIONADOS": reentries_by_year[next_year],
            "VALIDACION_CONGELADA": "OK",
            "ESTADO_PUBLICACION": "PUBLICABLE",
            "DESCRIPCION": (
                f"La tasa de rotación {base_year}-{next_year} fue "
                f"{format_value(rate, '%')}, "
                + (
                    f"{'disminuyendo' if variation < 0 else 'aumentando'} "
                    f"{format_value(abs(variation), 'pp')} respecto del periodo anterior. "
                    if variation is not None and variation != 0
                    else "sin comparación previa. "
                )
                + f"Permanecieron {remain} personas y rotaron {rotate}."
            ),
        })
        previous_rate = rate

    kpi_rows: list[dict[str, Any]] = []
    dictionary: dict[str, dict[str, Any]] = {}

    def add_kpi(
        code_name: str,
        indicator: str,
        dimension: str,
        period: Any,
        value: float | int | None,
        unit: str,
        formula: str,
        source: str,
        level: str,
        publicability: str,
        observation: str,
        sense: str = "INFORMATIVO",
        target: float | int | None = None,
        rule: str = "INFORMATIVO",
    ) -> None:
        kpi_rows.append({
            "CODIGO_KPI": code_name,
            "INDICADOR": indicator,
            "DIMENSION": dimension,
            "ANIO_PERIODO": period,
            "VALOR": value,
            "UNIDAD": unit,
            "SENTIDO": sense,
            "META": target,
            "FORMULA_METODO": formula,
            "FUENTE": source,
            "NIVEL_ANALISIS": level,
            "PUBLICABILIDAD": publicability,
            "OBSERVACION": observation,
            "_RULE": rule,
        })
        dictionary.setdefault(code_name, {
            "CODIGO_KPI": code_name,
            "INDICADOR": indicator,
            "DIMENSION": dimension,
            "DEFINICION": observation or indicator,
            "FORMULA": formula,
            "UNIDAD": unit,
            "SENTIDO": sense,
            "META_REFERENCIAL": target,
            "SEMAFORO": rule,
            "NIVEL_ANALISIS": level,
            "PUBLICABILIDAD": publicability,
            "LIMITACION": "Indicador agregado; no permite inferencia individual.",
            "FUENTE": source,
            "VERSION": "4.0",
        })

    for annual in annual_summary:
        year = annual["ANIO"]
        add_kpi("EST_DOTACION_TOTAL", "DOTACION_TOTAL", "ESTABILIDAD", year, annual["DOTACION_TOTAL"], "personas", "conteo de personas únicas por año", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Universo anual normalizado.")
        add_kpi("EST_REINGRESOS", "REINGRESOS", "ESTABILIDAD", year, annual["REINGRESOS"], "personas", "personas que reaparecen luego de un año ausente", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Indicador longitudinal descriptivo.")
        add_kpi("DIV_PORC_MUJERES", "PORC_MUJERES", "DIVERSIDAD", year, annual["PORC_MUJERES"], "%", "mujeres / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "No establece causalidad.", rule="INFORMATIVO")
        add_kpi("DIV_PORC_HOMBRES", "PORC_HOMBRES", "DIVERSIDAD", year, annual["PORC_HOMBRES"], "%", "hombres / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "No establece causalidad.", rule="INFORMATIVO")
        add_kpi("DIV_BRECHA_GENERO_PP", "BRECHA_GENERO_PP", "DIVERSIDAD", year, annual["BRECHA_GENERO_PP"], "pp", "abs(% hombres - % mujeres)", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Brecha descriptiva, no causal.", "BAJA_ES_FAVORABLE", 10.0, "BRECHA_GENERO")
        add_kpi("DIV_INDICE_PARIDAD_GENERO", "INDICE_PARIDAD_GENERO", "DIVERSIDAD", year, annual["INDICE_PARIDAD_GENERO"], "%", "min(hombres,mujeres) / max(hombres,mujeres) * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Índice descriptivo.", "SUBE_ES_FAVORABLE")
        add_kpi("DIV_EDAD_PROMEDIO", "EDAD_PROMEDIO", "DIVERSIDAD", year, annual["EDAD_PROMEDIO"], "años", "promedio edad al 31 de diciembre", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Edad calculada por año.")
        add_kpi("DIV_EDAD_MEDIANA", "EDAD_MEDIANA", "DIVERSIDAD", year, annual["EDAD_MEDIANA"], "años", "mediana edad al 31 de diciembre", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Edad calculada por año.")
        add_kpi("DIV_PORC_MENORES_30", "PORC_MENORES_30", "DIVERSIDAD", year, annual["PORC_MENORES_30"], "%", "personas menores de 30 / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Tramo etario descriptivo.")
        add_kpi("DIV_PORC_60_MAS", "PORC_60_MAS", "DIVERSIDAD", year, annual["PORC_60_MAS"], "%", "personas de 60 o más / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Tramo etario descriptivo.")
        add_kpi("CAP_HORAS_TOTALES", "HORAS_TOTALES", "CAPACIDAD", year, annual["HORAS_TOTALES"], "horas", "planta + contrata + honorarios", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "No equivale automáticamente a JCE.")
        add_kpi("CAP_HORAS_PROMEDIO", "HORAS_PROMEDIO", "CAPACIDAD", year, annual["HORAS_PROMEDIO"], "horas", "horas totales / personas", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Promedio descriptivo.")
        add_kpi("CAP_HORAS_MEDIANA", "HORAS_MEDIANA", "CAPACIDAD", year, annual["HORAS_MEDIANA"], "horas", "mediana horas contractuales", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Mediana descriptiva.")
        add_kpi("CAP_PORC_HORAS_PLANTA", "PORC_HORAS_PLANTA", "CAPACIDAD", year, annual["PORC_HORAS_PLANTA"], "%", "horas planta / horas totales * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Participación de componente horario.")
        add_kpi("CAP_PORC_HORAS_HONORARIOS", "PORC_HORAS_HONORARIOS", "CAPACIDAD", year, annual["PORC_HORAS_HONORARIOS"], "%", "horas honorarios / horas totales * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE", "Participación de componente horario.")
        add_kpi("CAP_PORC_PERSONAS_MAS_44", "PORC_PERSONAS_MAS_44", "CAPACIDAD", year, annual["PORC_MAS_44_HORAS"], "%", "personas con más de 44 horas / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Control descriptivo de horas contractuales.", "BAJA_ES_FAVORABLE", 0.0, "MAS_44")
        add_kpi("CAP_PERSONAS_MAS_44", "PERSONAS_MAS_44", "CAPACIDAD", year, annual["PERSONAS_MAS_44_HORAS"], "personas", "conteo personas con más de 44 horas contractuales", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Control descriptivo de horas contractuales.", "BAJA_ES_FAVORABLE")
        add_kpi("FOR_PORC_DOCTORADO", "PORC_DOCTORADO", "FORMACION", year, annual["PORC_DOCTORADO"], "%", "doctorado / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Categoría homologada por código oficial.")
        add_kpi("FOR_PORC_MAGISTER", "PORC_MAGISTER", "FORMACION", year, annual["PORC_MAGISTER"], "%", "magíster / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Categoría homologada por código oficial.")
        add_kpi("FOR_PORC_TITULO_PROFESIONAL", "PORC_TITULO_PROFESIONAL", "FORMACION", year, annual["PORC_TITULO_PROFESIONAL"], "%", "título profesional / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Categoría homologada por código oficial.")
        add_kpi("FOR_PORC_TECNICO_NIVEL_SUPERIOR", "PORC_TECNICO_NIVEL_SUPERIOR", "FORMACION", year, annual["PORC_TECNICO_NIVEL_SUPERIOR"], "%", "técnico nivel superior / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Categoría homologada por código oficial.")
        add_kpi("FOR_PORC_FORMACION_POSTGRADO", "PORC_FORMACION_POSTGRADO", "FORMACION", year, annual["PORC_FORMACION_POSTGRADO"], "%", "(doctorado + magíster) / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Agrupa códigos oficiales 1 y 2.")
        add_kpi("CAR_PORC_DOCENTES", "PORC_DOCENTES", "CARGO", year, annual["PORC_DOCENTES"], "%", "docentes / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Cargo homologado por código oficial.")
        add_kpi("CAR_PORC_DIRECTIVOS", "PORC_DIRECTIVOS", "CARGO", year, annual["PORC_DIRECTIVOS"], "%", "cargos directivos / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Agrupación descriptiva de cargos oficiales.")
        add_kpi("CAR_PORC_COORDINADORES", "PORC_COORDINADORES", "CARGO", year, annual["PORC_COORDINADORES"], "%", "coordinadores / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Cargo homologado por código oficial.")
        add_kpi("PRO_NUM_PROGRAMAS", "NUM_PROGRAMAS", "PROGRAMAS", year, annual["NUM_PROGRAMAS"], "programas", "conteo de programas homologados", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Homologación textual conservadora.")
        add_kpi("PRO_CONCENTRACION_TOP_1", "CONCENTRACION_TOP_1", "PROGRAMAS", year, annual["CONCENTRACION_TOP_1"], "%", "programa principal con más personas / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Concentración descriptiva.")
        add_kpi("PRO_CONCENTRACION_TOP_3", "CONCENTRACION_TOP_3", "PROGRAMAS", year, annual["CONCENTRACION_TOP_3"], "%", "tres programas principales / dotación total * 100", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "Concentración descriptiva.")
        add_kpi("PRO_INDICE_HHI_PROGRAMAS", "INDICE_HHI_PROGRAMAS", "PROGRAMAS", year, annual["INDICE_HHI_PROGRAMAS"], "índice", "suma de cuadrados de participaciones por programa", "08_DATOS_CRUZADOS", "ANUAL", "PUBLICABLE_CON_ADVERTENCIA", "HHI no es juicio de calidad por sí mismo.", "BAJA_ES_FAVORABLE", None, "HHI")

    quality_lookup = {
        (row["ANIO"], row["CONTROL"]): row
        for row in quality_rows
    }

    for annual in annual_summary:
        year = annual["ANIO"]
        quality_kpis = [
            ("CAL_COBERTURA_CLAVE", "COBERTURA_CLAVE", "clave documental completa"),
            ("CAL_COBERTURA_SEXO", "COBERTURA_SEXO", "sexo interpretable"),
            ("CAL_COBERTURA_EDAD", "COBERTURA_EDAD", "edad calculable"),
            ("CAL_COBERTURA_HORAS", "COBERTURA_HORAS", "horas totales calculables"),
            ("CAL_COBERTURA_FORMACION", "COBERTURA_FORMACION", "formación homologable"),
            ("CAL_COBERTURA_CARGO", "COBERTURA_CARGO", "cargo homologable"),
            ("CAL_COBERTURA_PROGRAMA", "COBERTURA_PROGRAMA", "programa informado"),
        ]

        for code_name, indicator, control in quality_kpis:
            row = quality_lookup[(year, control)]
            add_kpi(
                code_name,
                indicator,
                "CALIDAD",
                year,
                row["COBERTURA_PCT"],
                "%",
                f"conformes del control '{control}' / universo * 100",
                "07_CALIDAD_DATOS",
                "ANUAL",
                "PUBLICABLE",
                "Cobertura agregada sin PII.",
                "SUBE_ES_FAVORABLE",
                quality_targets[control],
                "COBERTURA",
            )

    for row in stability_rows:
        period = row["PERIODO"]
        add_kpi("EST_TASA_ROTACION", "TASA_ROTACION", "ESTABILIDAD", period, row["TASA_ROTACION"], "%", "rotan / dotación base * 100", "03_ESTABILIDAD", "PERIODO", "PUBLICABLE", "KPI oficial congelado; nuevos no participan en la tasa.", "BAJA_ES_FAVORABLE", None, "ROTACION")
        add_kpi("EST_TASA_PERMANENCIA", "TASA_PERMANENCIA", "ESTABILIDAD", period, row["TASA_PERMANENCIA"], "%", "permanecen / dotación base * 100", "03_ESTABILIDAD", "PERIODO", "PUBLICABLE", "Complemento descriptivo de permanencia.", "SUBE_ES_FAVORABLE")
        add_kpi("EST_PERMANECEN", "PERMANECEN", "ESTABILIDAD", period, row["PERMANECEN"], "personas", "intersección de personas entre base y siguiente", "03_ESTABILIDAD", "PERIODO", "PUBLICABLE", "Conteo agregado.")
        add_kpi("EST_ROTAN", "ROTAN", "ESTABILIDAD", period, row["ROTAN"], "personas", "personas de la base ausentes al año siguiente", "03_ESTABILIDAD", "PERIODO", "PUBLICABLE", "No equivale necesariamente a renuncia definitiva.", "BAJA_ES_FAVORABLE")
        add_kpi("EST_NUEVOS", "NUEVOS", "ESTABILIDAD", period, row["NUEVOS"], "personas", "personas nuevas en el año siguiente", "03_ESTABILIDAD", "PERIODO", "PUBLICABLE", "Los nuevos no participan en la tasa de rotación.")

    rows_by_code: dict[str, list[dict[str, Any]]] = {}

    for row in kpi_rows:
        rows_by_code.setdefault(row["CODIGO_KPI"], []).append(row)

    for rows in rows_by_code.values():
        rows.sort(key=lambda row: period_order(row["ANIO_PERIODO"]))
        values_so_far = []

        for index, row in enumerate(rows):
            previous = rows[index - 1]["VALOR"] if index else None
            value = row["VALOR"]
            gap = (
                value - previous
                if isinstance(value, (int, float))
                and isinstance(previous, (int, float))
                else None
            )
            variation_pct = (
                gap / abs(previous) * 100
                if gap is not None and previous not in (None, 0)
                else None
            )
            values_so_far.append(value)
            row["VALOR_ANTERIOR"] = previous
            row["BRECHA_ABSOLUTA"] = gap
            row["VARIACION_PCT"] = variation_pct
            row["TENDENCIA"] = trend_label(values_so_far)

            if row["_RULE"] == "ROTACION":
                row["SEMAFORO"] = rotation_light(gap)
            elif row["_RULE"] == "BRECHA_GENERO":
                row["SEMAFORO"] = gender_gap_light(value)
            elif row["_RULE"] == "MAS_44":
                row["SEMAFORO"] = above_44_light(value)
            elif row["_RULE"] == "COBERTURA":
                row["SEMAFORO"] = coverage_light(value, row["META"])
            elif row["_RULE"] == "HHI":
                row["SEMAFORO"] = hhi_light(value)
            else:
                row["SEMAFORO"] = "GRIS" if previous is None else "AMARILLO"

            row["BRECHA_META"] = (
                value - row["META"]
                if isinstance(value, (int, float))
                and isinstance(row["META"], (int, float))
                else None
            )
            row["DESCRIPCION_DINAMICA"] = description_text(
                row["INDICADOR"],
                row["ANIO_PERIODO"],
                value,
                row["UNIDAD"],
                previous,
                row["OBSERVACION"],
            )

    for row in kpi_rows:
        row.pop("_RULE", None)

    kpi_headers = [
        "CODIGO_KPI",
        "INDICADOR",
        "DIMENSION",
        "ANIO_PERIODO",
        "VALOR",
        "UNIDAD",
        "VALOR_ANTERIOR",
        "BRECHA_ABSOLUTA",
        "VARIACION_PCT",
        "TENDENCIA",
        "SENTIDO",
        "META",
        "BRECHA_META",
        "SEMAFORO",
        "DESCRIPCION_DINAMICA",
        "FORMULA_METODO",
        "FUENTE",
        "NIVEL_ANALISIS",
        "PUBLICABILIDAD",
        "OBSERVACION",
    ]
    stability_headers = [
        "PERIODO",
        "DOTACION_BASE",
        "DOTACION_SIGUIENTE",
        "PERMANECEN",
        "ROTAN",
        "NUEVOS",
        "TASA_ROTACION",
        "TASA_PERMANENCIA",
        "VARIACION_ROTACION_PP",
        "TENDENCIA_ROTACION",
        "SEMAFORO_ROTACION",
        "REINGRESOS_RELACIONADOS",
        "VALIDACION_CONGELADA",
        "ESTADO_PUBLICACION",
        "DESCRIPCION",
    ]
    diversity_headers = [
        "ANIO",
        "DOTACION_TOTAL",
        "HOMBRES",
        "MUJERES",
        "SIN_INFORMACION_SEXO",
        "PORC_HOMBRES",
        "PORC_MUJERES",
        "BRECHA_GENERO_PP",
        "INDICE_PARIDAD_GENERO",
        "EDAD_PROMEDIO",
        "EDAD_MEDIANA",
        "EDAD_MINIMA",
        "EDAD_MAXIMA",
        "PORC_MENORES_30",
        "PORC_30_39",
        "PORC_40_49",
        "PORC_50_59",
        "PORC_60_MAS",
        "VARIACION_PORC_MUJERES_PP",
        "VARIACION_EDAD_PROMEDIO",
        "SEMAFORO_BRECHA_GENERO",
        "DESCRIPCION",
    ]
    hours_headers = [
        "ANIO",
        "PERSONAS",
        "HORAS_TOTALES",
        "HORAS_PLANTA",
        "HORAS_CONTRATA",
        "HORAS_HONORARIOS",
        "HORAS_PROMEDIO",
        "HORAS_MEDIANA",
        "HORAS_MINIMA",
        "HORAS_MAXIMA",
        "PERSONAS_PLANTA",
        "PERSONAS_CONTRATA",
        "PERSONAS_HONORARIOS",
        "PERSONAS_COMBINADAS",
        "PERSONAS_SIN_HORAS",
        "PERSONAS_NO_DETERMINABLES",
        "PERSONAS_MAS_44_HORAS",
        "PORC_MAS_44_HORAS",
        "PORC_HORAS_PLANTA",
        "PORC_HORAS_HONORARIOS",
        "VARIACION_HORAS_TOTALES",
        "VARIACION_HORAS_PROMEDIO",
        "SEMAFORO_MAS_44",
        "DESCRIPCION",
    ]
    category_headers = [
        "ANIO",
        "DIMENSION",
        "CATEGORIA",
        "PERSONAS",
        "PORCENTAJE",
        "VARIACION_PERSONAS",
        "VARIACION_PORCENTAJE_PP",
        "TENDENCIA",
        "SEMAFORO",
        "DESCRIPCION",
    ]
    quality_headers = [
        "ANIO",
        "CONTROL",
        "UNIVERSO",
        "CONFORMES",
        "NO_CONFORMES",
        "COBERTURA_PCT",
        "BRECHA_META_PP",
        "ESTADO",
        "SEMAFORO",
        "DESCRIPCION",
        "FUENTE_VALIDACION",
    ]
    dictionary_headers = [
        "CODIGO_KPI",
        "INDICADOR",
        "DIMENSION",
        "DEFINICION",
        "FORMULA",
        "UNIDAD",
        "SENTIDO",
        "META_REFERENCIAL",
        "SEMAFORO",
        "NIVEL_ANALISIS",
        "PUBLICABILIDAD",
        "LIMITACION",
        "FUENTE",
        "VERSION",
    ]
    parameter_headers = [
        "PARAMETRO",
        "DIMENSION",
        "VALOR",
        "UNIDAD",
        "SENTIDO",
        "DESCRIPCION",
        "FUENTE",
        "VERSION",
    ]
    parameter_rows = [
        {"PARAMETRO": "UMBRAL_ROTACION_VERDE", "DIMENSION": "ESTABILIDAD", "VALOR": -2.0, "UNIDAD": "pp", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "Disminución superior a 2,0 pp.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "UMBRAL_ROTACION_ROJO", "DIMENSION": "ESTABILIDAD", "VALOR": 2.0, "UNIDAD": "pp", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "Aumento superior a 2,0 pp.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "META_COBERTURA_CRITICA", "DIMENSION": "CALIDAD", "VALOR": 100.0, "UNIDAD": "%", "SENTIDO": "SUBE_ES_FAVORABLE", "DESCRIPCION": "Meta para clave, duplicados, edad y horas.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "META_COBERTURA_GENERAL", "DIMENSION": "CALIDAD", "VALOR": 95.0, "UNIDAD": "%", "SENTIDO": "SUBE_ES_FAVORABLE", "DESCRIPCION": "Meta para sexo, formación, cargo y programa.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "TRAMOS_EDAD", "DIMENSION": "DIVERSIDAD", "VALOR": "<30;30-39;40-49;50-59;>=60", "UNIDAD": "años", "SENTIDO": "INFORMATIVO", "DESCRIPCION": "Tramos etarios usados.", "FUENTE": "HITO 2B-2", "VERSION": "4.0"},
        {"PARAMETRO": "TRAMOS_HORAS", "DIMENSION": "CAPACIDAD", "VALOR": "<11;11-22;23-43;44+", "UNIDAD": "horas", "SENTIDO": "INFORMATIVO", "DESCRIPCION": "Tramos horarios usados.", "FUENTE": "HITO 2B-2", "VERSION": "4.0"},
        {"PARAMETRO": "UMBRAL_44_HORAS", "DIMENSION": "CAPACIDAD", "VALOR": 44.0, "UNIDAD": "horas", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "Control descriptivo de personas sobre 44 horas.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "HHI_BAJA_CONCENTRACION", "DIMENSION": "PROGRAMAS", "VALOR": 0.15, "UNIDAD": "índice", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "HHI inferior a 0,15.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "HHI_ALTA_CONCENTRACION", "DIMENSION": "PROGRAMAS", "VALOR": 0.25, "UNIDAD": "índice", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "HHI superior a 0,25.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "TOLERANCIA_HORAS", "DIMENSION": "CAPACIDAD", "VALOR": 0.1, "UNIDAD": "horas", "SENTIDO": "CONTROL", "DESCRIPCION": "Tolerancia de validación de horas.", "FUENTE": "HITO 2B-2", "VERSION": "4.0"},
        {"PARAMETRO": "FECHA_CALCULO_EDAD", "DIMENSION": "DIVERSIDAD", "VALOR": "31-12 del año analizado", "UNIDAD": "fecha", "SENTIDO": "CONTROL", "DESCRIPCION": "Edad calculada por año, nunca con fecha actual.", "FUENTE": "HITO 2B-2", "VERSION": "4.0"},
        {"PARAMETRO": "VERSION_SIIPA", "DIMENSION": "GENERAL", "VALOR": "4.0", "UNIDAD": "versión", "SENTIDO": "CONTROL", "DESCRIPCION": "Versión objetivo del generador.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "REGLA_ROTACION", "DIMENSION": "ESTABILIDAD", "VALOR": "ROTAN / DOTACION_BASE * 100", "UNIDAD": "regla", "SENTIDO": "BAJA_ES_FAVORABLE", "DESCRIPCION": "Los nuevos no participan en la tasa.", "FUENTE": "Release KPI rotación v1.0", "VERSION": "4.0"},
        {"PARAMETRO": "REGLA_REINGRESO", "DIMENSION": "ESTABILIDAD", "VALOR": "persona reaparece tras ausencia intermedia", "UNIDAD": "regla", "SENTIDO": "INFORMATIVO", "DESCRIPCION": "Regla longitudinal descriptiva.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
        {"PARAMETRO": "PUBLICABILIDAD", "DIMENSION": "GENERAL", "VALOR": "PUBLICABLE;PUBLICABLE_CON_ADVERTENCIA;INTERNO", "UNIDAD": "categorías", "SENTIDO": "CONTROL", "DESCRIPCION": "Clasificación de uso por sensibilidad.", "FUENTE": "HITO 2B-3", "VERSION": "4.0"},
    ]

    append_table("02_KPI_MAESTRO", kpi_headers, kpi_rows)
    append_table("03_ESTABILIDAD", stability_headers, stability_rows)
    append_table("04_DIVERSIDAD", diversity_headers, diversity_rows)
    append_table("05_CAPACIDAD_HORAS", hours_headers, hours_rows)
    append_table("06_FORMACION_CARGO", category_headers, category_rows)
    append_table("07_CALIDAD_DATOS", quality_headers, quality_rows)
    append_table("90_DICCIONARIO_KPI", dictionary_headers, list(dictionary.values()))
    append_table("91_PARAMETROS", parameter_headers, parameter_rows)
    workbook["91_PARAMETROS"].sheet_state = "hidden"

    calculations_sheet = workbook["92_CALCULOS"]
    calculations_sheet.delete_rows(1, calculations_sheet.max_row)
    append_calculation_block(calculations_sheet, "RESUMEN_ANUAL", list(annual_summary[0].keys()), annual_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_ROTACION", stability_headers, stability_rows)
    append_calculation_block(calculations_sheet, "RESUMEN_GENERO", list(gender_summary[0].keys()), gender_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_EDAD", list(age_summary[0].keys()), age_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_HORAS", hours_headers, hours_rows)
    append_calculation_block(calculations_sheet, "RESUMEN_FORMACION", category_headers, formation_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_CARGO", category_headers, cargo_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_CONTRACTUAL", category_headers, contract_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_PROGRAMAS", category_headers, program_summary)
    append_calculation_block(calculations_sheet, "RESUMEN_CALIDAD", quality_headers, quality_rows)
    append_calculation_block(
        calculations_sheet,
        "REINGRESOS_AGREGADOS",
        ["ANIO", "REINGRESOS"],
        [
            {"ANIO": year, "REINGRESOS": value}
            for year, value in reentries_by_year.items()
        ],
    )
    append_calculation_block(
        calculations_sheet,
        "MOVILIDAD_AGREGADA",
        ["CONTROL", "VALOR", "OBSERVACION"],
        [{
            "CONTROL": "MOVILIDAD",
            "VALOR": None,
            "OBSERVACION": "No disponible en el generador para HITO 2B-3.",
        }],
    )

    def category_count(
        year: int,
        dimension: str,
        category: str,
    ) -> int:
        for row in category_rows:
            if (
                row["ANIO"] == year
                and row["DIMENSION"] == dimension
                and row["CATEGORIA"] == category
            ):
                return row["PERSONAS"]

        return 0

    append_chart_block(
        "CHART_DOTACION",
        ["ANIO", "DOTACION_TOTAL"],
        [
            {
                "ANIO": row["ANIO"],
                "DOTACION_TOTAL": row["DOTACION_TOTAL"],
            }
            for row in annual_summary
        ],
    )
    append_chart_block(
        "CHART_ROTACION",
        ["PERIODO", "TASA_ROTACION"],
        [
            {
                "PERIODO": row["PERIODO"],
                "TASA_ROTACION": row["TASA_ROTACION"],
            }
            for row in stability_rows
        ],
    )
    append_chart_block(
        "CHART_MOVIMIENTOS",
        ["PERIODO", "PERMANECEN", "ROTAN", "NUEVOS"],
        [
            {
                "PERIODO": row["PERIODO"],
                "PERMANECEN": row["PERMANECEN"],
                "ROTAN": row["ROTAN"],
                "NUEVOS": row["NUEVOS"],
            }
            for row in stability_rows
        ],
    )
    append_chart_block(
        "CHART_SEXO",
        ["ANIO", "HOMBRES", "MUJERES"],
        gender_summary,
    )
    append_chart_block(
        "CHART_TRAMOS_EDAD",
        ["ANIO", "<30", "30-39", "40-49", "50-59", ">=60"],
        [
            {
                "ANIO": year,
                "<30": category_count(year, "TRAMO_EDAD", "<30"),
                "30-39": category_count(year, "TRAMO_EDAD", "30-39"),
                "40-49": category_count(year, "TRAMO_EDAD", "40-49"),
                "50-59": category_count(year, "TRAMO_EDAD", "50-59"),
                ">=60": category_count(year, "TRAMO_EDAD", ">=60"),
            }
            for year in EXPECTED_UNIVERSES
        ],
    )
    append_chart_block(
        "CHART_HORAS",
        ["ANIO", "HORAS_TOTALES", "HORAS_PROMEDIO"],
        [
            {
                "ANIO": row["ANIO"],
                "HORAS_TOTALES": row["HORAS_TOTALES"],
                "HORAS_PROMEDIO": row["HORAS_PROMEDIO"],
            }
            for row in annual_summary
        ],
    )
    append_chart_block(
        "CHART_CONTRACTUAL",
        ["ANIO", "PLANTA", "CONTRATA", "HONORARIOS", "COMBINADO", "SIN_HORAS", "NO_DETERMINABLE"],
        [
            {
                "ANIO": year,
                "PLANTA": category_count(year, "TIPO_CONTRACTUAL", "PLANTA"),
                "CONTRATA": category_count(year, "TIPO_CONTRACTUAL", "CONTRATA"),
                "HONORARIOS": category_count(year, "TIPO_CONTRACTUAL", "HONORARIOS"),
                "COMBINADO": category_count(year, "TIPO_CONTRACTUAL", "COMBINADO"),
                "SIN_HORAS": category_count(year, "TIPO_CONTRACTUAL", "SIN_HORAS"),
                "NO_DETERMINABLE": category_count(year, "TIPO_CONTRACTUAL", "NO_DETERMINABLE"),
            }
            for year in EXPECTED_UNIVERSES
        ],
    )
    append_chart_block(
        "CHART_FORMACION",
        ["ANIO", "DOCTORADO", "MAGISTER", "TITULO PROFESIONAL", "OTRA"],
        [
            {
                "ANIO": year,
                "DOCTORADO": category_count(year, "FORMACION", "DOCTORADO"),
                "MAGISTER": category_count(year, "FORMACION", "MAGISTER"),
                "TITULO PROFESIONAL": category_count(year, "FORMACION", "TITULO PROFESIONAL"),
                "OTRA": sum(
                    row["PERSONAS"]
                    for row in category_rows
                    if row["ANIO"] == year
                    and row["DIMENSION"] == "FORMACION"
                    and row["CATEGORIA"] not in {
                        "DOCTORADO",
                        "MAGISTER",
                        "TITULO PROFESIONAL",
                    }
                ),
            }
            for year in EXPECTED_UNIVERSES
        ],
    )
    append_chart_block(
        "CHART_CARGO_DOCENTE",
        ["ANIO", "DOCENTES", "OTROS_CARGOS"],
        [
            {
                "ANIO": row["ANIO"],
                "DOCENTES": category_count(row["ANIO"], "CARGO", "DOCENTE"),
                "OTROS_CARGOS": row["DOTACION_TOTAL"] - category_count(row["ANIO"], "CARGO", "DOCENTE"),
            }
            for row in annual_summary
        ],
    )
    append_chart_block(
        "CHART_CALIDAD",
        ["ANIO", "CLAVE", "SEXO", "EDAD", "HORAS", "FORMACION", "CARGO", "PROGRAMA"],
        [
            {
                "ANIO": year,
                "CLAVE": quality_lookup[(year, "clave documental completa")]["COBERTURA_PCT"],
                "SEXO": quality_lookup[(year, "sexo interpretable")]["COBERTURA_PCT"],
                "EDAD": quality_lookup[(year, "edad calculable")]["COBERTURA_PCT"],
                "HORAS": quality_lookup[(year, "horas totales calculables")]["COBERTURA_PCT"],
                "FORMACION": quality_lookup[(year, "formación homologable")]["COBERTURA_PCT"],
                "CARGO": quality_lookup[(year, "cargo homologable")]["COBERTURA_PCT"],
                "PROGRAMA": quality_lookup[(year, "programa informado")]["COBERTURA_PCT"],
            }
            for year in EXPECTED_UNIVERSES
        ],
    )
    calculations_sheet.sheet_state = "hidden"

    for sheet_name in (
        "02_KPI_MAESTRO",
        "03_ESTABILIDAD",
        "04_DIVERSIDAD",
        "05_CAPACIDAD_HORAS",
        "06_FORMACION_CARGO",
        "07_CALIDAD_DATOS",
        "90_DICCIONARIO_KPI",
        "91_PARAMETROS",
    ):
        format_numeric_columns(sheet_name)

    build_cover()
    build_dashboard()

    chart_inventory = [
        add_chart("Dotación anual 2022-2026", "line", "CHART_DOTACION", [2], "A35"),
        add_chart("Tasa de rotación por periodo", "line", "CHART_ROTACION", [2], "J35"),
        add_chart("Permanecen, rotan y nuevos", "col", "CHART_MOVIMIENTOS", [2, 3, 4], "A50"),
        add_chart("Hombres y mujeres por año", "col", "CHART_SEXO", [2, 3], "J50"),
        add_chart("Distribución por tramo de edad", "stacked", "CHART_TRAMOS_EDAD", [2, 3, 4, 5, 6], "A65"),
        add_chart("Horas totales y promedio", "line", "CHART_HORAS", [2, 3], "J65"),
        add_chart("Distribución contractual por año", "stacked", "CHART_CONTRACTUAL", [2, 3, 4, 5, 6, 7], "A80"),
        add_chart("Formación académica por año", "stacked", "CHART_FORMACION", [2, 3, 4, 5], "J80"),
        add_chart("Docentes versus otros cargos", "col", "CHART_CARGO_DOCENTE", [2, 3], "A95"),
        add_chart("Cobertura de calidad de datos", "line", "CHART_CALIDAD", [2, 3, 4, 5, 6, 7, 8], "J95"),
    ]

    for sheet_name in (
        "02_KPI_MAESTRO",
        "03_ESTABILIDAD",
        "04_DIVERSIDAD",
        "05_CAPACIDAD_HORAS",
        "06_FORMACION_CARGO",
        "07_CALIDAD_DATOS",
        "08_DATOS_CRUZADOS",
        "90_DICCIONARIO_KPI",
        "91_PARAMETROS",
    ):
        apply_number_formats(sheet_name)
        apply_visual_states(sheet_name)
        apply_common_table_format(
            sheet_name,
            maximum_width=30 if sheet_name == "08_DATOS_CRUZADOS" else 42,
        )

    for year in EXPECTED_UNIVERSES:
        raw_name = worksheet_for_year(year)
        apply_common_table_format(
            raw_name,
            maximum_width=30,
        )

    workbook["91_PARAMETROS"].sheet_state = "hidden"
    workbook["92_CALCULOS"].sheet_state = "hidden"
    workbook["93_VALIDACIONES"].sheet_state = "visible"

    allowed_lights = {"VERDE", "AMARILLO", "ROJO", "GRIS"}

    validation_rows = [
        validation_row(
            "numero_hojas",
            len(SIIPA_SHEET_NAMES),
            len(workbook.sheetnames),
            len(workbook.sheetnames) == len(SIIPA_SHEET_NAMES),
            "Conteo de hojas del workbook generado.",
        ),
        validation_row(
            "orden_hojas",
            " | ".join(SIIPA_SHEET_NAMES),
            " | ".join(workbook.sheetnames),
            workbook.sheetnames == SIIPA_SHEET_NAMES,
            "Orden exacto exigido para SIIPA.",
        ),
        validation_row(
            "persona_anio",
            sum(EXPECTED_UNIVERSES.values()),
            total_person_year,
            total_person_year == sum(EXPECTED_UNIVERSES.values()),
            "Registros homologados persona-año en 08_DATOS_CRUZADOS.",
        ),
        validation_row(
            "duplicados_persona_anio",
            0,
            duplicate_person_year,
            duplicate_person_year == 0,
            "Duplicados por ANIO + CLAVE_PERSONA.",
        ),
        validation_row(
            "claves_incompletas",
            0,
            incomplete_keys,
            incomplete_keys == 0,
            "Claves incompletas en datos cruzados.",
        ),
        validation_row(
            "formulas_raw_datos_cruzados",
            0,
            formulas_in_data,
            formulas_in_data == 0,
            "Fórmulas detectadas en hojas RAW o datos cruzados.",
        ),
        validation_row(
            "referencia_v3_1",
            0,
            sum(
                1
                for value in used_paths
                if "v3.1" in value.lower()
            ),
            not any(
                "v3.1" in value.lower()
                for value in used_paths
            ),
            "Las rutas efectivamente usadas no apuntan a SIIPA v3.1.",
        ),
        validation_row(
            "acceso_onedrive",
            0,
            sum(
                1
                for value in used_paths
                if "onedrive" in value.lower()
            ),
            not any(
                "onedrive" in value.lower()
                for value in used_paths
            ),
            "Las fuentes leídas pertenecen a normalized_restricted local.",
        ),
    ]

    for year, source in sources.items():
        records = [
            record
            for record in crossed_records
            if record["ANIO"] == year
        ]
        sexes = Counter(
            record["SEXO_NORMALIZADO"]
            for record in records
        )
        ages = [
            record["EDAD"]
            for record in records
            if record["EDAD"] is not None
        ]
        hour_values = [
            record["TOTAL_HORAS_CONTRACTUALES"]
            for record in records
            if record["TOTAL_HORAS_CONTRACTUALES"] is not None
        ]
        raw_sheet = workbook[worksheet_for_year(year)]

        validation_rows.extend([
            validation_row(
                f"universo_{year}",
                EXPECTED_UNIVERSES[year],
                len(source.rows),
                len(source.rows) == EXPECTED_UNIVERSES[year],
                f"Filas del CSV normalizado {source.path.name}.",
            ),
            validation_row(
                f"raw_{year}_encabezados",
                "encabezados_reales",
                raw_sheet.max_column,
                raw_sheet.max_column == len(source.headers)
                and [
                    raw_sheet.cell(1, index + 1).value
                    for index in range(len(source.headers))
                ] == source.headers,
                "La fila 1 de RAW replica los encabezados del CSV.",
            ),
            validation_row(
                f"datos_cruzados_{year}",
                EXPECTED_UNIVERSES[year],
                len(records),
                len(records) == EXPECTED_UNIVERSES[year],
                "Registros homologados por año.",
            ),
            validation_row(
                f"edad_calculable_{year}",
                EXPECTED_UNIVERSES[year],
                len(ages),
                len(ages) == EXPECTED_UNIVERSES[year],
                "Edad calculada al 31 de diciembre del año analizado.",
            ),
            validation_row(
                f"sexo_{year}_hombre",
                expected_sex[year]["HOMBRE"],
                sexes.get("HOMBRE", 0),
                sexes.get("HOMBRE", 0) == expected_sex[year]["HOMBRE"],
                "Conteo anual de hombres.",
            ),
            validation_row(
                f"sexo_{year}_mujer",
                expected_sex[year]["MUJER"],
                sexes.get("MUJER", 0),
                sexes.get("MUJER", 0) == expected_sex[year]["MUJER"],
                "Conteo anual de mujeres.",
            ),
            validation_row(
                f"sexo_{year}_identidad",
                EXPECTED_UNIVERSES[year],
                sexes.get("HOMBRE", 0) + sexes.get("MUJER", 0),
                (
                    sexes.get("HOMBRE", 0)
                    + sexes.get("MUJER", 0)
                ) == EXPECTED_UNIVERSES[year],
                "HOMBRE + MUJER debe igualar el universo anual.",
            ),
            validation_row(
                f"horas_calculables_{year}",
                EXPECTED_UNIVERSES[year],
                len(hour_values),
                len(hour_values) == EXPECTED_UNIVERSES[year],
                "Total contractual calculado desde planta, contrata y honorarios.",
            ),
            validation_row(
                f"horas_total_{year}",
                expected_hours[year],
                round(sum(hour_values), 2),
                abs(sum(hour_values) - expected_hours[year]) <= 0.1,
                "Suma de horas contractuales con tolerancia 0,1.",
            ),
        ])

    for row in stability_rows:
        validation_rows.extend([
            validation_row(
                f"rotacion_congelada_{row['PERIODO']}",
                "OK",
                row["VALIDACION_CONGELADA"],
                row["VALIDACION_CONGELADA"] == "OK",
                "KPI de rotación coincide con la línea base congelada.",
            ),
            validation_row(
                f"rotacion_tasa_{row['PERIODO']}",
                round(row["ROTAN"] / row["DOTACION_BASE"] * 100, 6),
                round(row["TASA_ROTACION"], 6),
                abs(
                    row["ROTAN"] / row["DOTACION_BASE"] * 100
                    - row["TASA_ROTACION"]
                ) <= 0.000001,
                "Tasa calculada como ROTAN / DOTACION_BASE * 100.",
            ),
            validation_row(
                f"rotacion_identidad_base_{row['PERIODO']}",
                row["DOTACION_BASE"],
                row["PERMANECEN"] + row["ROTAN"],
                row["PERMANECEN"] + row["ROTAN"] == row["DOTACION_BASE"],
                "PERMANECEN + ROTAN = DOTACION_BASE.",
            ),
            validation_row(
                f"rotacion_identidad_siguiente_{row['PERIODO']}",
                row["DOTACION_SIGUIENTE"],
                row["PERMANECEN"] + row["NUEVOS"],
                row["PERMANECEN"] + row["NUEVOS"] == row["DOTACION_SIGUIENTE"],
                "PERMANECEN + NUEVOS = DOTACION_SIGUIENTE.",
            ),
        ])

    for annual in annual_summary:
        year = annual["ANIO"]
        total = annual["DOTACION_TOTAL"]
        year_categories = [
            row
            for row in category_rows
            if row["ANIO"] == year
        ]

        for dimension in (
            "FORMACION",
            "CARGO",
            "TIPO_CONTRACTUAL",
            "TRAMO_EDAD",
            "TRAMO_HORAS",
            "PROGRAMA",
        ):
            rows = [
                row
                for row in year_categories
                if row["DIMENSION"] == dimension
            ]
            people_sum = sum(row["PERSONAS"] for row in rows)
            percent_sum = sum(row["PORCENTAJE"] for row in rows)
            validation_rows.extend([
                validation_row(
                    f"suma_{dimension.lower()}_{year}",
                    total,
                    people_sum,
                    people_sum == total,
                    f"Suma de personas por {dimension}.",
                ),
                validation_row(
                    f"porcentaje_{dimension.lower()}_{year}",
                    100.0,
                    round(percent_sum, 4),
                    abs(percent_sum - 100.0) <= 0.1,
                    f"Suma porcentual por {dimension}.",
                ),
            ])

        validation_rows.extend([
            validation_row(
                f"suma_genero_{year}",
                total,
                annual["HOMBRES"] + annual["MUJERES"] + annual["SIN_INFORMACION_SEXO"],
                annual["HOMBRES"] + annual["MUJERES"] + annual["SIN_INFORMACION_SEXO"] == total,
                "Hombres + mujeres + sin información = universo.",
            ),
            validation_row(
                f"horas_componentes_{year}",
                round(annual["HORAS_TOTALES"], 2),
                round(
                    annual["HORAS_PLANTA"]
                    + annual["HORAS_CONTRATA"]
                    + annual["HORAS_HONORARIOS"],
                    2,
                ),
                abs(
                    annual["HORAS_TOTALES"]
                    - (
                        annual["HORAS_PLANTA"]
                        + annual["HORAS_CONTRATA"]
                        + annual["HORAS_HONORARIOS"]
                    )
                ) <= 0.1,
                "Horas totales igualan suma de componentes.",
            ),
        ])

    kpi_key_pairs = [
        (
            row["CODIGO_KPI"],
            row["ANIO_PERIODO"],
        )
        for row in kpi_rows
    ]
    kpi_codes = {
        row["CODIGO_KPI"]
        for row in kpi_rows
    }
    dictionary_codes = set(dictionary)
    semaphores = [
        row["SEMAFORO"]
        for row in kpi_rows
    ] + [
        row["SEMAFORO_ROTACION"]
        for row in stability_rows
    ] + [
        row["SEMAFORO_BRECHA_GENERO"]
        for row in diversity_rows
    ] + [
        row["SEMAFORO_MAS_44"]
        for row in hours_rows
    ] + [
        row["SEMAFORO"]
        for row in category_rows
    ] + [
        row["SEMAFORO"]
        for row in quality_rows
    ]
    empty_descriptions = sum(
        1
        for row in kpi_rows
        if not text(row.get("DESCRIPCION_DINAMICA"))
    )
    analytic_sheets = [
        "02_KPI_MAESTRO",
        "03_ESTABILIDAD",
        "04_DIVERSIDAD",
        "05_CAPACIDAD_HORAS",
        "06_FORMACION_CARGO",
        "07_CALIDAD_DATOS",
        "90_DICCIONARIO_KPI",
        "91_PARAMETROS",
        "92_CALCULOS",
    ]
    pii_tokens = {
        "TIPO_DOCUMENTO",
        "NUM_DOCUMENTO",
        "PRIMER_APELLIDO",
        "SEGUNDO_APELLIDO",
        "NOMBRES",
        "CLAVE_PERSONA",
    }
    pii_hits = 0

    for sheet_name in analytic_sheets:
        sheet = workbook[sheet_name]

        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in pii_tokens:
                    pii_hits += 1

    validation_rows.extend([
        validation_row(
            "kpi_maestro_poblado",
            "mayor que 0",
            len(kpi_rows),
            len(kpi_rows) > 0,
            "02_KPI_MAESTRO contiene indicadores calculados.",
        ),
        validation_row(
            "kpi_maestro_sin_duplicados",
            len(kpi_key_pairs),
            len(set(kpi_key_pairs)),
            len(kpi_key_pairs) == len(set(kpi_key_pairs)),
            "Clave CODIGO_KPI + ANIO_PERIODO sin duplicados.",
        ),
        validation_row(
            "diccionario_completo",
            len(kpi_codes),
            len(kpi_codes & dictionary_codes),
            kpi_codes.issubset(dictionary_codes),
            "Todo KPI del maestro tiene definición.",
        ),
        validation_row(
            "diccionario_sin_sobrantes",
            len(dictionary_codes),
            len(kpi_codes & dictionary_codes),
            dictionary_codes.issubset(kpi_codes),
            "El diccionario sólo contiene KPI del maestro.",
        ),
        validation_row(
            "semaforos_textuales",
            "VERDE/AMARILLO/ROJO/GRIS",
            "; ".join(sorted(set(semaphores))),
            set(semaphores).issubset(allowed_lights),
            "Semáforos restringidos a estados textuales permitidos.",
        ),
        validation_row(
            "descripciones_dinamicas",
            0,
            empty_descriptions,
            empty_descriptions == 0,
            "Todos los KPI del maestro tienen descripción generada por Python.",
        ),
        validation_row(
            "formulas_excel_total",
            0,
            sum(has_formula(sheet) for sheet in workbook.worksheets),
            sum(has_formula(sheet) for sheet in workbook.worksheets) == 0,
            "El workbook no contiene fórmulas Excel.",
        ),
        validation_row(
            "nan_inf_total",
            0,
            bad_number_count(SIIPA_SHEET_NAMES),
            bad_number_count(SIIPA_SHEET_NAMES) == 0,
            "No hay valores NaN o infinitos.",
        ),
        validation_row(
            "pii_hojas_analiticas",
            0,
            pii_hits,
            pii_hits == 0,
            "Hojas analíticas no contienen columnas de identificación individual.",
        ),
        validation_row(
            "parametros_documentados",
            "mayor que 0",
            len(parameter_rows),
            len(parameter_rows) > 0,
            "91_PARAMETROS documenta reglas y umbrales usados.",
        ),
        validation_row(
            "calculos_auxiliares",
            "mayor que 0",
            workbook["92_CALCULOS"].max_row,
            workbook["92_CALCULOS"].max_row > 1,
            "92_CALCULOS contiene resultados auxiliares calculados en Python.",
        ),
    ])

    dashboard_charts = workbook["01_RESUMEN_EJECUTIVO"]._charts
    chart_series_ok = all(
        len(chart.series) > 0
        for chart in dashboard_charts
    )
    chart_categories_ok = all(
        all(getattr(series, "cat", None) is not None for series in chart.series)
        for chart in dashboard_charts
    )
    chart_refs = " ".join(
        text(row["RANGO_CATEGORIAS"]) + " " + text(row["RANGO_SERIES"])
        for row in chart_inventory
    )
    public_analytic_sheets = [
        sheet_name
        for sheet_name in SIIPA_SHEET_NAMES
        if not re.match(r"^(0[9]|1[0-3])_RAW_", sheet_name)
        and sheet_name != "08_DATOS_CRUZADOS"
    ]
    workbook_strings = []

    for sheet_name in public_analytic_sheets:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    workbook_strings.append(cell.value)

    validation_rows.extend([
        validation_row(
            "portada",
            "poblada",
            workbook["00_PORTADA"]["A1"].value,
            workbook["00_PORTADA"]["A1"].value == "SISTEMA INTEGRADO DE INDICADORES DE PERSONAL ACADÉMICO",
            "Portada institucional generada.",
        ),
        validation_row(
            "dashboard_ejecutivo",
            "poblado",
            workbook["01_RESUMEN_EJECUTIVO"]["A1"].value,
            workbook["01_RESUMEN_EJECUTIVO"]["A1"].value == "Resumen ejecutivo SIIPA v4.0",
            "Resumen ejecutivo con tarjetas y síntesis.",
        ),
        validation_row(
            "graficos_total",
            10,
            len(dashboard_charts),
            len(dashboard_charts) == 10,
            "Gráficos ubicados en 01_RESUMEN_EJECUTIVO.",
        ),
        validation_row(
            "graficos_series",
            "series no vacías",
            chart_series_ok,
            chart_series_ok,
            "Cada gráfico contiene al menos una serie.",
        ),
        validation_row(
            "graficos_categorias",
            "categorías no vacías",
            chart_categories_ok,
            chart_categories_ok,
            "Cada gráfico contiene categorías.",
        ),
        validation_row(
            "graficos_sin_ref",
            0,
            chart_refs.count("#REF!"),
            "#REF!" not in chart_refs,
            "Rangos de gráficos sin referencias rotas.",
        ),
        validation_row(
            "parametros_ocultos",
            "hidden",
            workbook["91_PARAMETROS"].sheet_state,
            workbook["91_PARAMETROS"].sheet_state == "hidden",
            "91_PARAMETROS oculta, no veryHidden.",
        ),
        validation_row(
            "calculos_ocultos",
            "hidden",
            workbook["92_CALCULOS"].sheet_state,
            workbook["92_CALCULOS"].sheet_state == "hidden",
            "92_CALCULOS oculta, no veryHidden.",
        ),
        validation_row(
            "validaciones_visibles",
            "visible",
            workbook["93_VALIDACIONES"].sheet_state,
            workbook["93_VALIDACIONES"].sheet_state == "visible",
            "93_VALIDACIONES visible para revisión técnica.",
        ),
        validation_row(
            "enlaces_externos",
            0,
            len(getattr(workbook, "_external_links", [])),
            len(getattr(workbook, "_external_links", [])) == 0,
            "No existen vínculos externos en el workbook.",
        ),
        validation_row(
            "rutas_cloud_operativas",
            0,
            sum(1 for value in workbook_strings if "onedrive" in value.lower()),
            not any("onedrive" in value.lower() for value in workbook_strings),
            "No existen rutas cloud operativas en hojas públicas o analíticas.",
        ),
        validation_row(
            "texto_centrado_muestra",
            "center",
            workbook["02_KPI_MAESTRO"]["A1"].alignment.horizontal,
            workbook["02_KPI_MAESTRO"]["A1"].alignment.horizontal == "center",
            "Encabezados y celdas analíticas centradas.",
        ),
        validation_row(
            "decimales_muestra",
            "0.0",
            workbook["02_KPI_MAESTRO"]["E2"].number_format,
            workbook["02_KPI_MAESTRO"]["E2"].number_format in {"0.0", "0"},
            "Valores numéricos persistidos con formato decimal según tipo.",
        ),
    ])

    validation_sheet = workbook["93_VALIDACIONES"]
    validation_sheet.delete_rows(1, validation_sheet.max_row)
    validation_sheet.append([
        "CONTROL",
        "ESPERADO",
        "OBSERVADO",
        "ESTADO",
        "DETALLE",
    ])

    for row in validation_rows:
        validation_sheet.append([
            row["CONTROL"],
            row["ESPERADO"],
            row["OBSERVADO"],
            row["ESTADO"],
            row["DETALLE"],
        ])

    apply_number_formats("93_VALIDACIONES")
    apply_visual_states("93_VALIDACIONES")
    apply_common_table_format("93_VALIDACIONES", maximum_width=42)
    workbook["93_VALIDACIONES"].sheet_state = "visible"

    workbook.save(output)

    for data_only in (False, True):
        reopened = load_workbook(
            output,
            read_only=False,
            data_only=data_only,
        )

        try:
            if len(reopened.sheetnames) != len(SIIPA_SHEET_NAMES):
                raise RuntimeError(
                    f"Workbook inválido: {len(reopened.sheetnames)} hojas."
                )

            if reopened.sheetnames != SIIPA_SHEET_NAMES:
                raise RuntimeError(
                    "Workbook inválido: orden o nombres de hojas incorrectos."
                )

            if (
                reopened["08_DATOS_CRUZADOS"].max_row - 1
                != sum(EXPECTED_UNIVERSES.values())
            ):
                raise RuntimeError(
                    "Workbook inválido: 08_DATOS_CRUZADOS no contiene 552 filas."
                )

            formulas_reopened = sum(
                has_formula(sheet)
                for sheet in reopened.worksheets
            )

            if formulas_reopened != 0:
                raise RuntimeError(
                    "Workbook inválido: contiene fórmulas Excel."
                )

            literal_errors_reopened = sum(
                literal_errors(sheet)
                for sheet in reopened.worksheets
            )

            if literal_errors_reopened != 0:
                raise RuntimeError(
                    "Workbook inválido: errores literales de fórmula."
                )

            for year, source in sources.items():
                sheet = reopened[worksheet_for_year(year)]

                if sheet.max_row - 1 != EXPECTED_UNIVERSES[year]:
                    raise RuntimeError(
                        f"Workbook inválido: RAW {year} no tiene "
                        f"{EXPECTED_UNIVERSES[year]} registros."
                    )

                headers = [
                    sheet.cell(1, column_index + 1).value
                    for column_index in range(len(source.headers))
                ]

                if headers != source.headers:
                    raise RuntimeError(
                        f"Workbook inválido: encabezados RAW {year}."
                    )

            if reopened["02_KPI_MAESTRO"].max_row <= 1:
                raise RuntimeError(
                    "Workbook inválido: KPI maestro no poblado."
                )

            if reopened["90_DICCIONARIO_KPI"].max_row <= 1:
                raise RuntimeError(
                    "Workbook inválido: diccionario KPI no poblado."
                )

            if reopened["91_PARAMETROS"].max_row <= 1:
                raise RuntimeError(
                    "Workbook inválido: parámetros no poblados."
                )

            if reopened["92_CALCULOS"].max_row <= 1:
                raise RuntimeError(
                    "Workbook inválido: cálculos auxiliares no poblados."
                )
        finally:
            reopened.close()

    if any(
        row["ESTADO"] != "OK"
        for row in validation_rows
    ):
        failed = [
            row["CONTROL"]
            for row in validation_rows
            if row["ESTADO"] != "OK"
        ]

        raise RuntimeError(
            "Exportación SIIPA con validaciones en revisión: "
            + ", ".join(failed)
        )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Generador oficial del SIIPA de Personal Académico."
        )
    )

    parser.add_argument(
        "--diagnostico",
        action="store_true",
        help=(
            "Valida fuentes, universos, rotación, "
            "reingresos, edad, sexo y horas."
        ),
    )

    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help="Ruta del Excel final.",
    )

    parser.add_argument(
        "--exportar",
        action="store_true",
        help="Genera el Excel final.",
    )

    return parser


def main() -> int:
    args = build_parser().parse_args()

    if args.diagnostico:
        result = diagnostic()

        print(
            json.dumps(
                result,
                ensure_ascii=False,
                indent=2,
            )
        )

        print()
        print("GENERADOR_DIAGNOSTICO_OK")

    if args.exportar:
        export_excel(args.output)

    if not args.diagnostico and not args.exportar:
        print(
            "Use --diagnostico o --exportar.",
            file=sys.stderr,
        )
        return 2

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
