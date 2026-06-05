#!/usr/bin/env python3
"""Incorpora la base preliminar editable de Personal Academico SIES 2026.

Lee el Excel controlado con openpyxl, genera RAW TSV, normaliza a la
estructura oficial de Personal Academico en la Institucion y mantiene una base
general acumulable. No genera archivo final PES_READY.
"""

from __future__ import annotations

import csv
from collections import Counter, defaultdict
from copy import copy
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
import unicodedata

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
MODULE_DIR = SCRIPT_DIR.parent
REPO_DIR = MODULE_DIR.parent

EXCEL_ORIGINAL = Path("/Users/alexi/Desktop/Libro1.xlsx")
EXCEL_COPIA = MODULE_DIR / "insumos/preliminares/Libro1_base_preliminar_personal_academico_2026.xlsx"
EXCEL_BACKUPS_DIR = MODULE_DIR / "insumos/preliminares/backups"
RAW_TSV = MODULE_DIR / "insumos/preliminares/base_preliminar_personal_academico_en_institucion_RAW.tsv"
MAPEO_TSV = MODULE_DIR / "catalogos/mapeo_encabezados_base_preliminar_en_institucion.tsv"
BASE_NORMALIZADA = MODULE_DIR / "data/base_general/base_preliminar_en_institucion_normalizada.tsv"
BASE_GENERAL = MODULE_DIR / "data/base_general/base_general_personal_academico_en_institucion.tsv"
BASE_GENERAL_BACKUPS_DIR = MODULE_DIR / "data/base_general/backups"
AUDITORIAS_DIR = MODULE_DIR / "auditorias/base_preliminar"
REPORTE_MD = MODULE_DIR / "docs/REPORTE_FASE_3_BASE_PRELIMINAR.md"
BITACORA_MD = MODULE_DIR / "bitacoras/base_preliminar/BITACORA_FASE_3_BASE_PRELIMINAR.md"
REPORTE_3B_MD = MODULE_DIR / "docs/REPORTE_FASE_3B_DEPURACION_FECHAS_Y_NA.md"
BITACORA_3B_MD = MODULE_DIR / "bitacoras/base_preliminar/BITACORA_FASE_3B_DEPURACION_FECHAS_Y_NA.md"
RESUMEN_JSON = MODULE_DIR / "auditorias/base_preliminar/resumen_ultima_incorporacion_base_preliminar.json"
FECHA_REFERENCIA = date(2026, 6, 11)

COLUMNAS_OFICIALES = [
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "NIVEL_FORMACION_ACADEMICO",
    "NOMBRE_TITULO_O_GRADO",
    "NOMBRE_INSTITUCION_OBT_TITULO",
    "PAIS_OBTENCION_TIT_O_GRADO",
    "FECHA_OBT_TIT_O_GRADO",
    "NIVEL_FORMACION_ESPECIALIDAD",
    "TIPO_ESPECIALIDAD",
    "NOMBRE_ESPECIALIDAD",
    "NOMBRE_INST_OBT_ESPECIALIDAD",
    "PAIS_OBTENCION_ESPECIALIDAD",
    "FECHA_OBTENCION_ESPECIALIDAD",
    "PRINCIPAL_CARGO_ACADEMICO",
    "CARGO_NORMALIZADO",
    "NIVEL_SUPERIOR_ADSCRIPCION",
    "NIVEL_SECUNDARIO_ADSCRIPCION",
    "COMUNA_MAYOR_FUNCION",
    "NOMBRE_PRINCIPAL_PROGRAMA",
    "TOTAL_HORAS_PRINCIPAL_PROGRAMA",
    "COMUNA_PRINCIPAL_PROGRAMA",
    "NUM_HORAS_PLANTA",
    "NUM_HORAS_CONTRATA",
    "NUM_HORAS_HONORARIOS",
    "JERARQUIA_ACADEMICA",
    "JERARQUIA_ACADEMICA_OCDE",
    "VIGENCIA",
]

MAPEO_OBLIGATORIO = {
    "TIPO_DOCUMENTO": "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO": "NUM_DOCUMENTO",
    "DV": "DV",
    "PRIMER_APELLIDO": "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO": "SEGUNDO_APELLIDO",
    "NOMBRES": "NOMBRES",
    "SEXO": "SEXO",
    "FECHA_NACIMIENTO": "FECHA_NACIMIENTO",
    "NACIONALIDAD": "NACIONALIDAD",
    "NIVEL_FORMACIÓN_ACADÉM": "NIVEL_FORMACION_ACADEMICO",
    "NOMBRE_TÍTULO O GRADO": "NOMBRE_TITULO_O_GRADO",
    "NOMB_INSTIT_OBT_TÍTULO": "NOMBRE_INSTITUCION_OBT_TITULO",
    "PAÍS_OBT_TÍTULO_O_GRADO": "PAIS_OBTENCION_TIT_O_GRADO",
    "FECHA_OBT_TÍTULO_O_GRADO": "FECHA_OBT_TIT_O_GRADO",
    "NIVEL_FORMACIÓN_ESPECIALIDAD": "NIVEL_FORMACION_ESPECIALIDAD",
    "TIPO_ESPECIALIDAD": "TIPO_ESPECIALIDAD",
    "NOMBRE_ESPECIALIDAD": "NOMBRE_ESPECIALIDAD",
    "NOMBRE_INSTITUC_OBT_ESPECIALIDAD": "NOMBRE_INST_OBT_ESPECIALIDAD",
    "PAÍS_OBT_ESPECIALIDAD": "PAIS_OBTENCION_ESPECIALIDAD",
    "FECHA_OBT_ESPECIALIDAD": "FECHA_OBTENCION_ESPECIALIDAD",
    "PRINCIPAL_CARGO_ACADÉMICO": "PRINCIPAL_CARGO_ACADEMICO",
    "CARGO_NORMALIZADO": "CARGO_NORMALIZADO",
    "NIVEL_SUPERIOR_ADSCRIPCIÓN": "NIVEL_SUPERIOR_ADSCRIPCION",
    "NIVEL_SECUNDARIO_ADSCRIPCIÓN": "NIVEL_SECUNDARIO_ADSCRIPCION",
    "COMUNA_MAYOR_FUNCIÓN": "COMUNA_MAYOR_FUNCION",
    "NOMBRE_PRINCIPAL_PROGRAMA": "NOMBRE_PRINCIPAL_PROGRAMA",
    "TOTAL_HORAS_PRINCIPAL_PROGRAMA": "TOTAL_HORAS_PRINCIPAL_PROGRAMA",
    "COMUNA_PRINCIPAL_PROGRAMA": "COMUNA_PRINCIPAL_PROGRAMA",
    "NUM_HORAS_PLANTA": "NUM_HORAS_PLANTA",
    "NUM_HORAS_CONTRATA": "NUM_HORAS_CONTRATA",
    "NUM_HORAS_HONORARIO": "NUM_HORAS_HONORARIOS",
    "JERARQUÍA_ACADÉMICA": "JERARQUIA_ACADEMICA",
    "JERARQUÍA_ACADÉMICA_OCDE": "JERARQUIA_ACADEMICA_OCDE",
    "VIGENCIA": "VIGENCIA",
}

FUENTE_ESTRUCTURA = "ESTRUCTURA_OFICIAL"
FUENTE_NORMATIVA = "REGLA_NORMATIVA"
FUENTE_PREVENTIVA = "VALIDACION_TECNICA_PREVENTIVA"
FUENTE_PENDIENTE = "PENDIENTE_CONFIRMACION"

COLUMNAS_FECHA = {
    "FECHA_NACIMIENTO",
    "FECHA_OBT_TIT_O_GRADO",
    "FECHA_OBTENCION_ESPECIALIDAD",
}
COLUMNAS_HORAS = {
    "TOTAL_HORAS_PRINCIPAL_PROGRAMA",
    "NUM_HORAS_PLANTA",
    "NUM_HORAS_CONTRATA",
    "NUM_HORAS_HONORARIOS",
}
COLUMNAS_OBLIGATORIAS_PREVENTIVAS = {
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "NIVEL_FORMACION_ACADEMICO",
    "VIGENCIA",
}
COLUMNAS_MAYUSCULAS = set(COLUMNAS_OFICIALES)
VALORES_NO_DISPONIBLES = {"#N/A", "#N/D", "N/A", "NA", "NO APLICA"}
CLAVE_TECNICA = [
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "NOMBRE_PRINCIPAL_PROGRAMA",
    "PRINCIPAL_CARGO_ACADEMICO",
]


class IncorporacionError(RuntimeError):
    """Error critico que debe detener la incorporacion."""


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def asegurar_directorios() -> None:
    for path in [
        EXCEL_COPIA.parent,
        EXCEL_BACKUPS_DIR,
        RAW_TSV.parent,
        MAPEO_TSV.parent,
        BASE_NORMALIZADA.parent,
        BASE_GENERAL_BACKUPS_DIR,
        AUDITORIAS_DIR,
        REPORTE_MD.parent,
        BITACORA_MD.parent,
    ]:
        path.mkdir(parents=True, exist_ok=True)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(REPO_DIR))
    except ValueError:
        return str(path)


def limpiar_invisibles(texto: str) -> str:
    return "".join(
        ch for ch in texto if unicodedata.category(ch) not in {"Cf", "Cc"} or ch in "\t\n\r"
    )


def clave_encabezado(texto: object) -> str:
    valor = "" if texto is None else str(texto)
    valor = limpiar_invisibles(valor)
    valor = re.sub(r"\s+", " ", valor).strip()
    valor = unicodedata.normalize("NFD", valor)
    valor = "".join(ch for ch in valor if unicodedata.category(ch) != "Mn")
    valor = unicodedata.normalize("NFC", valor).upper()
    valor = valor.replace(" ", "_")
    valor = re.sub(r"[^A-Z0-9_Ñ]", "_", valor)
    valor = re.sub(r"_+", "_", valor).strip("_")
    return valor


def construir_mapeo_normalizado() -> dict[str, str]:
    mapeo = {}
    for preliminar, oficial in MAPEO_OBLIGATORIO.items():
        mapeo[clave_encabezado(preliminar)] = oficial
    for oficial in COLUMNAS_OFICIALES:
        mapeo[clave_encabezado(oficial)] = oficial
    return mapeo


def serializar_celda(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    return str(valor)


def normalizar_disponibilidad(valor: str) -> str:
    return re.sub(r"\s+", " ", valor.strip()).upper()


def fila_vacia(fila: list[str]) -> bool:
    return all(celda == "" or celda.strip() == "" for celda in fila)


class Auditor:
    def __init__(self, path: Path, archivo_origen: Path) -> None:
        self.path = path
        self.archivo_origen = archivo_origen
        self.eventos: list[dict[str, object]] = []

    def add(
        self,
        tipo_evento: str,
        severidad: str,
        mensaje: str,
        fuente_regla: str,
        *,
        hoja_origen: str = "",
        fila_origen: object = "",
        columna_origen: str = "",
        columna_oficial: str = "",
        valor_original: object = "",
        valor_normalizado: object = "",
        accion: str = "",
    ) -> None:
        self.eventos.append(
            {
                "timestamp": datetime.now().isoformat(timespec="seconds"),
                "archivo_origen": str(self.archivo_origen),
                "hoja_origen": hoja_origen,
                "fila_origen": fila_origen,
                "tipo_evento": tipo_evento,
                "severidad": severidad,
                "columna_origen": columna_origen,
                "columna_oficial": columna_oficial,
                "valor_original": "" if valor_original is None else valor_original,
                "valor_normalizado": "" if valor_normalizado is None else valor_normalizado,
                "accion": accion,
                "mensaje": mensaje,
                "fuente_regla": fuente_regla,
            }
        )

    def escribir(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        campos = [
            "timestamp",
            "archivo_origen",
            "hoja_origen",
            "fila_origen",
            "tipo_evento",
            "severidad",
            "columna_origen",
            "columna_oficial",
            "valor_original",
            "valor_normalizado",
            "accion",
            "mensaje",
            "fuente_regla",
        ]
        with self.path.open("w", newline="", encoding="utf-8-sig") as fh:
            writer = csv.DictWriter(fh, fieldnames=campos)
            writer.writeheader()
            writer.writerows(self.eventos)


def copiar_excel_controlado(auditor: Auditor) -> tuple[str, str, str]:
    if not EXCEL_ORIGINAL.exists():
        raise IncorporacionError(f"No existe Excel original: {EXCEL_ORIGINAL}")
    backup = ""
    if EXCEL_COPIA.exists():
        backup_path = EXCEL_BACKUPS_DIR / (
            f"Libro1_base_preliminar_personal_academico_2026_{timestamp()}.xlsx"
        )
        shutil.copy2(EXCEL_COPIA, backup_path)
        backup = rel(backup_path)
    shutil.copy2(EXCEL_ORIGINAL, EXCEL_COPIA)
    hash_original = sha256(EXCEL_ORIGINAL)
    hash_copia = sha256(EXCEL_COPIA)
    auditor.add(
        "EXCEL_COPIADO",
        "INFO",
        f"Excel copiado a insumo controlado. Backup previo: {backup or 'no aplica'}.",
        FUENTE_PREVENTIVA,
        accion="COPIAR_EXCEL_CONTROLADO",
        valor_original=str(EXCEL_ORIGINAL),
        valor_normalizado=rel(EXCEL_COPIA),
    )
    auditor.add(
        "HASH_EXCEL_ORIGINAL",
        "INFO",
        "Hash SHA256 del Excel original.",
        FUENTE_PREVENTIVA,
        valor_original=str(EXCEL_ORIGINAL),
        valor_normalizado=hash_original,
    )
    auditor.add(
        "HASH_EXCEL_COPIA",
        "INFO",
        "Hash SHA256 de la copia controlada.",
        FUENTE_PREVENTIVA,
        valor_original=rel(EXCEL_COPIA),
        valor_normalizado=hash_copia,
    )
    return hash_original, hash_copia, backup


def auditar_hojas(wb, auditor: Auditor) -> list[dict[str, object]]:
    hojas = []
    for ws in wb.worksheets:
        filas = 0
        columnas = ws.max_column or 0
        vacias = 0
        celdas_con_valor = 0
        posibles_encabezados = []
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            valores = [serializar_celda(v) for v in row]
            if any(v.strip() for v in valores):
                filas += 1
                celdas_con_valor += sum(1 for v in valores if v.strip())
                claves = {clave_encabezado(v) for v in valores if str(v).strip()}
                if {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV", "PRIMER_APELLIDO", "VIGENCIA"}.issubset(
                    claves
                ):
                    posibles_encabezados.append(idx)
            else:
                vacias += 1
        info = {
            "nombre": ws.title,
            "filas_utiles": filas,
            "filas_excel": ws.max_row or 0,
            "columnas": columnas,
            "filas_vacias": vacias,
            "celdas_con_valor": celdas_con_valor,
            "posibles_encabezados": posibles_encabezados,
        }
        hojas.append(info)
        auditor.add(
            "HOJA_DETECTADA",
            "INFO",
            (
                f"Hoja detectada con {filas} filas utiles, {columnas} columnas, "
                f"{vacias} filas completamente vacias y {celdas_con_valor} celdas con valores."
            ),
            FUENTE_PREVENTIVA,
            hoja_origen=ws.title,
            valor_normalizado=json.dumps(info, ensure_ascii=False),
        )
    return hojas


def seleccionar_hoja(hojas: list[dict[str, object]], auditor: Auditor) -> str:
    con_datos = [h for h in hojas if int(h["filas_utiles"]) > 0]
    if not con_datos:
        raise IncorporacionError("No hay hojas con datos en el Excel.")
    if len(con_datos) == 1:
        seleccionada = con_datos[0]
        razon = "Unica hoja con datos."
    else:
        candidatas = [h for h in con_datos if int(h["columnas"]) >= 34]
        if not candidatas:
            raise IncorporacionError("Hay varias hojas, pero ninguna tiene al menos 34 columnas.")
        seleccionada = sorted(
            candidatas,
            key=lambda h: (int(h["filas_utiles"]), int(h["celdas_con_valor"])),
            reverse=True,
        )[0]
        razon = "Mayor numero de filas utiles entre hojas con al menos 34 columnas potenciales."
    auditor.add(
        "HOJA_SELECCIONADA",
        "INFO",
        razon,
        FUENTE_PREVENTIVA,
        hoja_origen=str(seleccionada["nombre"]),
        valor_normalizado=json.dumps(seleccionada, ensure_ascii=False),
        accion="SELECCIONAR_HOJA_PRINCIPAL",
    )
    return str(seleccionada["nombre"])


def detectar_encabezado(ws, auditor: Auditor) -> tuple[int, list[str], list[int]]:
    requeridos = {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV", "PRIMER_APELLIDO", "VIGENCIA"}
    encabezados_posibles = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        valores = [serializar_celda(v) for v in row]
        claves = {clave_encabezado(v) for v in valores if str(v).strip()}
        if requeridos.issubset(claves):
            encabezados_posibles.append(idx)
    if not encabezados_posibles:
        raise IncorporacionError("No se encontro encabezado claro con columnas requeridas.")
    fila_encabezado = encabezados_posibles[0]
    encabezado = [
        serializar_celda(c.value)
        for c in next(ws.iter_rows(min_row=fila_encabezado, max_row=fila_encabezado))
    ]
    auditor.add(
        "ENCABEZADO_DETECTADO",
        "INFO",
        f"Encabezado detectado en fila {fila_encabezado}.",
        FUENTE_ESTRUCTURA,
        hoja_origen=ws.title,
        fila_origen=fila_encabezado,
        valor_normalizado=" | ".join(encabezado),
    )
    repetidos_internos = encabezados_posibles[1:]
    for fila in repetidos_internos:
        auditor.add(
            "ENCABEZADO_REPETIDO_INTERNO",
            "WARN",
            "Fila interna con patron de encabezado oficial detectada.",
            FUENTE_PREVENTIVA,
            hoja_origen=ws.title,
            fila_origen=fila,
        )
    return fila_encabezado, encabezado, repetidos_internos


def escribir_tsv(path: Path, encabezado: list[str], filas: list[list[str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh, delimiter="\t", lineterminator="\n")
        writer.writerow(encabezado)
        writer.writerows(filas)


def leer_tsv_dicts(path: Path) -> list[dict[str, str]]:
    with path.open("r", newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def generar_raw(ws, fila_encabezado: int, encabezado: list[str], auditor: Auditor) -> list[list[str]]:
    ancho = len(encabezado)
    filas = []
    for row in ws.iter_rows(min_row=fila_encabezado + 1, values_only=True):
        valores = [serializar_celda(v) for v in row]
        if len(valores) < ancho:
            valores += [""] * (ancho - len(valores))
        elif len(valores) > ancho:
            extras = valores[ancho:]
            if any(v.strip() for v in extras):
                auditor.add(
                    "RAW_FILA_ANCHO_INVALIDO",
                    "ERROR",
                    f"Fila con valores fuera del ancho RAW esperado de {ancho} columnas.",
                    FUENTE_ESTRUCTURA,
                    hoja_origen=ws.title,
                    fila_origen=len(filas) + fila_encabezado + 1,
                    valor_original=" | ".join(extras),
                )
            valores = valores[:ancho]
        if fila_vacia(valores):
            continue
        filas.append(valores)
    errores_ancho = [e for e in auditor.eventos if e["tipo_evento"] == "RAW_FILA_ANCHO_INVALIDO"]
    if errores_ancho:
        raise IncorporacionError("Existen filas RAW con ancho invalido que impiden construir TSV.")
    escribir_tsv(RAW_TSV, encabezado, filas)
    auditor.add(
        "RAW_TSV_CREADO",
        "INFO",
        f"RAW TSV creado con {len(filas)} registros de datos y {ancho} columnas.",
        FUENTE_ESTRUCTURA,
        hoja_origen=ws.title,
        valor_normalizado=rel(RAW_TSV),
        accion="CREAR_RAW_TSV",
    )
    return filas


def analizar_fecha(valor: str) -> tuple[str, date | None]:
    texto = valor.strip()
    if not texto:
        return "VACIA", None
    if normalizar_disponibilidad(texto) in VALORES_NO_DISPONIBLES:
        return "NO_DISPONIBLE", None
    patrones = [
        (r"^\d{1,2}[-/]\d{1,2}[-/]\d{2}$", "AMBIGUA"),
        (r"^\d{1,2}[-/]\d{1,2}[-/]\d{4}$", "CUATRO_DIGITOS"),
        (r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$", "CUATRO_DIGITOS"),
    ]
    for patron, estado in patrones:
        if re.match(patron, texto):
            for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y"):
                try:
                    return estado, datetime.strptime(texto, fmt).date()
                except ValueError:
                    pass
            return "INVALIDA", None
    return "INVALIDA", None


def validar_raw(
    encabezado: list[str],
    filas: list[list[str]],
    repetidos_internos: list[int],
    auditor: Auditor,
    hoja: str,
) -> dict[str, int]:
    metricas = Counter()
    metricas["total_columnas_raw"] = len(encabezado)
    metricas["total_filas_raw"] = len(filas) + 1
    metricas["total_filas_utiles"] = len(filas)
    metricas["total_encabezados_repetidos_internos"] = len(repetidos_internos)
    if len(encabezado) != 34:
        auditor.add(
            "RAW_FILA_ANCHO_INVALIDO",
            "ERROR",
            f"Encabezado RAW tiene {len(encabezado)} columnas; se esperaban 34.",
            FUENTE_ESTRUCTURA,
            hoja_origen=hoja,
        )
        raise IncorporacionError("El encabezado RAW no tiene 34 columnas.")
    for idx, fila in enumerate(filas, start=2):
        if len(fila) != 34:
            auditor.add(
                "RAW_FILA_ANCHO_INVALIDO",
                "ERROR",
                f"Fila RAW con {len(fila)} columnas; se esperaban 34.",
                FUENTE_ESTRUCTURA,
                hoja_origen=hoja,
                fila_origen=idx,
            )
            raise IncorporacionError("Existen filas RAW con ancho distinto de 34.")
        for col, valor in zip(encabezado, fila):
            disponible = normalizar_disponibilidad(valor)
            if disponible == "#N/A":
                metricas["total_valores_excel_na"] += 1
                auditor.add(
                    "VALOR_NO_DISPONIBLE_EXCEL",
                    "WARN",
                    "Valor #N/A leido desde Excel y preservado como dato fuente no disponible.",
                    FUENTE_PREVENTIVA,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_origen=col,
                    valor_original=valor,
                )
            elif disponible == "#N/D":
                metricas["total_valores_nd"] += 1
                auditor.add(
                    "VALOR_ND",
                    "WARN",
                    "Valor #N/D preservado para revision.",
                    FUENTE_PENDIENTE,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_origen=col,
                    valor_original=valor,
                )
            elif disponible in {"N/A", "NA"}:
                metricas["total_valores_na"] += 1
                auditor.add(
                    "VALOR_NO_DISPONIBLE_EXCEL",
                    "WARN",
                    "Valor N/A o NA preservado como dato fuente no disponible.",
                    FUENTE_PREVENTIVA,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_origen=col,
                    valor_original=valor,
                )
            elif disponible == "NO APLICA":
                metricas["total_valores_no_aplica"] += 1
            if valor.strip() == "":
                metricas[f"vacios::{col}"] += 1
    return dict(metricas)


def escribir_catalogo_mapeo(encabezado: list[str], auditor: Auditor, hoja: str) -> dict[int, str]:
    mapeo_norm = construir_mapeo_normalizado()
    asignaciones: dict[int, str] = {}
    filas_catalogo = []
    oficiales_asignadas = Counter()
    for idx, preliminar in enumerate(encabezado):
        oficial = mapeo_norm.get(clave_encabezado(preliminar), "")
        estado = "MAPEADO" if oficial else "NO_OFICIAL"
        observacion = (
            "Mapeo por estructura oficial y normalizacion tecnica preventiva de encabezado."
            if oficial
            else "Columna preliminar no coincide con estructura oficial."
        )
        if oficial:
            asignaciones[idx] = oficial
            oficiales_asignadas[oficial] += 1
            auditor.add(
                "MAPEO_ENCABEZADO",
                "INFO",
                observacion,
                FUENTE_ESTRUCTURA,
                hoja_origen=hoja,
                columna_origen=preliminar,
                columna_oficial=oficial,
                valor_original=preliminar,
                valor_normalizado=oficial,
            )
        else:
            auditor.add(
                "COLUMNA_PRELIMINAR_NO_OFICIAL",
                "WARN",
                observacion,
                FUENTE_ESTRUCTURA,
                hoja_origen=hoja,
                columna_origen=preliminar,
            )
        filas_catalogo.append(
            {
                "encabezado_preliminar": preliminar,
                "encabezado_normalizado": clave_encabezado(preliminar),
                "encabezado_oficial": oficial,
                "estado_mapeo": estado,
                "observacion": observacion,
            }
        )
    faltantes = [col for col in COLUMNAS_OFICIALES if oficiales_asignadas[col] == 0]
    duplicados = [col for col, total in oficiales_asignadas.items() if total > 1]
    for col in faltantes:
        auditor.add(
            "COLUMNA_OFICIAL_FALTANTE",
            "ERROR",
            "Columna oficial faltante en el mapeo preliminar.",
            FUENTE_ESTRUCTURA,
            hoja_origen=hoja,
            columna_oficial=col,
        )
    for col in duplicados:
        auditor.add(
            "MAPEO_ENCABEZADO",
            "ERROR",
            "Mas de una columna preliminar mapea a la misma columna oficial.",
            FUENTE_ESTRUCTURA,
            hoja_origen=hoja,
            columna_oficial=col,
        )
    if faltantes or duplicados or len(asignaciones) != 34:
        raise IncorporacionError("No se pudo mapear a 34 columnas oficiales de forma univoca.")
    with MAPEO_TSV.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(
            fh,
            delimiter="\t",
            fieldnames=[
                "encabezado_preliminar",
                "encabezado_normalizado",
                "encabezado_oficial",
                "estado_mapeo",
                "observacion",
            ],
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(filas_catalogo)
    return asignaciones


def normalizar_valor(valor: str, columna: str) -> tuple[str, bool]:
    original = valor
    texto = re.sub(r"\s+", " ", valor).strip()
    if texto.upper() == "NO APLICA":
        texto = "NO APLICA"
    elif columna == "DV" and texto == "k":
        texto = "K"
    elif columna in COLUMNAS_MAYUSCULAS and texto != "#N/D":
        texto = texto.upper()
    return texto, texto != original


def normalizar_base(
    encabezado: list[str],
    filas: list[list[str]],
    asignaciones: dict[int, str],
    auditor: Auditor,
    hoja: str,
) -> tuple[list[dict[str, str]], dict[str, int]]:
    normalizadas = []
    metricas = Counter()
    documentos = defaultdict(list)
    for idx, fila in enumerate(filas, start=2):
        claves_fila = {clave_encabezado(valor) for valor in fila if valor.strip()}
        if {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV", "PRIMER_APELLIDO", "VIGENCIA"}.issubset(
            claves_fila
        ):
            auditor.add(
                "ENCABEZADO_REPETIDO_INTERNO",
                "WARN",
                "Encabezado interno preservado en RAW y omitido de base normalizada.",
                FUENTE_PREVENTIVA,
                hoja_origen=hoja,
                fila_origen=idx,
                accion="OMITIR_DE_NORMALIZACION",
            )
            continue
        registro = {col: "" for col in COLUMNAS_OFICIALES}
        for pos, valor in enumerate(fila):
            oficial = asignaciones[pos]
            normalizado, cambio = normalizar_valor(valor, oficial)
            registro[oficial] = normalizado
            if cambio:
                evento = "DV_MINUSCULA_NORMALIZADA" if oficial == "DV" and valor == "k" else "TEXTO_NORMALIZADO"
                fuente = FUENTE_PREVENTIVA
                auditor.add(
                    evento,
                    "INFO",
                    "Valor normalizado con regla tecnica permitida.",
                    fuente,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_origen=encabezado[pos],
                    columna_oficial=oficial,
                    valor_original=valor,
                    valor_normalizado=normalizado,
                    accion="NORMALIZAR_VALOR",
                )
                if evento == "DV_MINUSCULA_NORMALIZADA":
                    metricas["total_dv_minuscula_normalizada"] += 1
        for col in COLUMNAS_OBLIGATORIAS_PREVENTIVAS:
            if registro[col].strip() == "":
                metricas["total_campos_obligatorios_preventivos_vacios"] += 1
                auditor.add(
                    "CAMPO_OBLIGATORIO_PREVENTIVO_VACIO",
                    "WARN",
                    "Campo vacio segun control tecnico preventivo.",
                    FUENTE_PREVENTIVA,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                )
        for col in COLUMNAS_FECHA:
            estado, fecha = analizar_fecha(registro[col])
            if estado == "AMBIGUA":
                metricas["total_fechas_ambiguas"] += 1
                auditor.add(
                    "FECHA_AMBIGUA",
                    "WARN",
                    "Fecha con anio de dos digitos o formato ambiguo; se preserva original.",
                    FUENTE_PENDIENTE,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=registro[col],
                )
            elif estado == "CUATRO_DIGITOS":
                metricas["total_fechas_anio_cuatro_digitos"] += 1
            elif estado == "NO_DISPONIBLE":
                metricas["total_campos_fecha_no_disponibles"] += 1
                auditor.add(
                    "VALOR_NO_DISPONIBLE_EXCEL",
                    "WARN",
                    "Campo de fecha con valor fuente no disponible preservado.",
                    FUENTE_PREVENTIVA,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=registro[col],
                )
            elif estado == "INVALIDA":
                metricas["total_fechas_invalidas"] += 1
                auditor.add(
                    "FECHA_INVALIDA",
                    "WARN",
                    "Fecha no interpretable; se preserva original.",
                    FUENTE_PENDIENTE,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=registro[col],
                )
            if fecha and fecha > FECHA_REFERENCIA:
                metricas["total_fechas_futuras"] += 1
                auditor.add(
                    "FECHA_FUTURA",
                    "WARN",
                    "Fecha futura respecto a 2026-06-11; se preserva original.",
                    FUENTE_PREVENTIVA,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=registro[col],
                )
        for col in COLUMNAS_HORAS:
            valor = registro[col]
            if "," in valor:
                metricas["total_horas_coma_decimal"] += 1
                auditor.add(
                    "HORA_COMA_DECIMAL",
                    "WARN",
                    "Hora con coma decimal preservada para revision.",
                    FUENTE_PENDIENTE,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=valor,
                )
            if valor and valor != "#N/D" and not re.match(r"^\d+([,.]\d+)?$", valor):
                metricas["total_horas_invalidas"] += 1
                auditor.add(
                    "HORA_INVALIDA",
                    "WARN",
                    "Hora no numerica preservada para revision.",
                    FUENTE_PENDIENTE,
                    hoja_origen=hoja,
                    fila_origen=idx,
                    columna_oficial=col,
                    valor_original=valor,
                )
        doc_key = (registro["TIPO_DOCUMENTO"], registro["NUM_DOCUMENTO"], registro["DV"])
        documentos[doc_key].append(
            (registro["NOMBRE_PRINCIPAL_PROGRAMA"], registro["PRINCIPAL_CARGO_ACADEMICO"], idx)
        )
        normalizadas.append(registro)
    for doc_key, apariciones in documentos.items():
        if len(apariciones) > 1:
            metricas["total_posibles_duplicados_documento"] += len(apariciones)
            combinaciones = {(prog, cargo) for prog, cargo, _ in apariciones}
            evento = "DOCUMENTO_MULTIPLES_FUNCIONES" if len(combinaciones) > 1 else "POSIBLE_DUPLICADO_DOCUMENTO"
            if evento == "DOCUMENTO_MULTIPLES_FUNCIONES":
                metricas["total_documentos_multiples_funciones"] += 1
            for prog, cargo, fila in apariciones:
                auditor.add(
                    evento,
                    "WARN",
                    "Documento aparece mas de una vez; no se elimina automaticamente.",
                    FUENTE_NORMATIVA,
                    hoja_origen=hoja,
                    fila_origen=fila,
                    valor_original="|".join(doc_key),
                    valor_normalizado=f"{prog}|{cargo}",
                )
    with BASE_NORMALIZADA.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, delimiter="\t", fieldnames=COLUMNAS_OFICIALES, lineterminator="\n")
        writer.writeheader()
        writer.writerows(normalizadas)
    return normalizadas, dict(metricas)


def clave_registro(registro: dict[str, str]) -> tuple[str, ...]:
    return tuple(registro.get(col, "") for col in CLAVE_TECNICA)


def escribir_base_general(registros: list[dict[str, str]], auditor: Auditor) -> dict[str, int | str]:
    metricas: dict[str, int | str] = {
        "total_filas_incorporadas": 0,
        "total_filas_actualizadas": 0,
        "base_general_backup": "",
    }
    existentes: list[dict[str, str]] = []
    if BASE_GENERAL.exists():
        backup_path = BASE_GENERAL_BACKUPS_DIR / (
            f"base_general_personal_academico_en_institucion_{timestamp()}.tsv"
        )
        shutil.copy2(BASE_GENERAL, backup_path)
        metricas["base_general_backup"] = rel(backup_path)
        auditor.add(
            "BASE_GENERAL_BACKUP",
            "INFO",
            "Backup creado antes de actualizar base general.",
            FUENTE_PREVENTIVA,
            valor_normalizado=rel(backup_path),
            accion="CREAR_BACKUP_BASE_GENERAL",
        )
        existentes = leer_tsv_dicts(BASE_GENERAL)
        depurados = []
        for reg in existentes:
            if (
                reg.get("TIPO_DOCUMENTO") == "TIPO_DOCUMENTO"
                and reg.get("NUM_DOCUMENTO") == "NUM_DOCUMENTO"
                and reg.get("DV") == "DV"
            ):
                auditor.add(
                    "ENCABEZADO_REPETIDO_INTERNO",
                    "WARN",
                    "Residuo de encabezado interno detectado en base general previa; se omite porque no es registro de persona.",
                    FUENTE_PREVENTIVA,
                    valor_original="|".join(clave_registro(reg)),
                    accion="OMITIR_ENCABEZADO_RESIDUAL_BASE_GENERAL",
                )
                continue
            depurados.append(reg)
        existentes = depurados
    indice = {clave_registro(reg): copy(reg) for reg in existentes}
    orden = [clave_registro(reg) for reg in existentes]
    for reg in registros:
        key = clave_registro(reg)
        if key in indice:
            indice[key] = reg
            metricas["total_filas_actualizadas"] = int(metricas["total_filas_actualizadas"]) + 1
            auditor.add(
                "FILA_ACTUALIZADA",
                "INFO",
                "Registro reemplazado por coincidencia de clave tecnica.",
                FUENTE_PREVENTIVA,
                valor_original="|".join(key),
                accion="ACTUALIZAR_BASE_GENERAL",
            )
        else:
            indice[key] = reg
            orden.append(key)
            metricas["total_filas_incorporadas"] = int(metricas["total_filas_incorporadas"]) + 1
            auditor.add(
                "FILA_INCORPORADA",
                "INFO",
                "Registro incorporado a base general acumulable.",
                FUENTE_PREVENTIVA,
                valor_original="|".join(key),
                accion="INCORPORAR_BASE_GENERAL",
            )
    salida = [indice[key] for key in orden]
    with BASE_GENERAL.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, delimiter="\t", fieldnames=COLUMNAS_OFICIALES, lineterminator="\n")
        writer.writeheader()
        writer.writerows(salida)
    auditor.add(
        "BASE_GENERAL_ACTUALIZADA" if existentes else "BASE_GENERAL_CREADA",
        "INFO",
        f"Base general {'actualizada' if existentes else 'creada'} con {len(salida)} filas.",
        FUENTE_PREVENTIVA,
        valor_normalizado=rel(BASE_GENERAL),
    )
    metricas["total_filas_base_general"] = len(salida)
    return metricas


def clasificar_fecha_diagnostico(valor: str) -> tuple[str, date | None, str]:
    estado, fecha = analizar_fecha(valor)
    if estado == "VACIA":
        return "FECHA_VACIA", None, "Campo de fecha vacio."
    if estado == "NO_DISPONIBLE":
        return (
            "FECHA_NO_DISPONIBLE",
            None,
            "Valor fuente no disponible preservado; no se interpreta como fecha.",
        )
    if estado == "AMBIGUA":
        return "FECHA_AMBIGUA", fecha, "Fecha interpretable con anio de dos digitos."
    if estado == "INVALIDA":
        return "FECHA_INVALIDA", None, "Formato no reconocido y distinto de valores no disponibles."
    if fecha and fecha > FECHA_REFERENCIA:
        return "FECHA_FUTURA", fecha, "Fecha posterior a 2026-06-11."
    return "FECHA_OK", fecha, "Fecha interpretable."


def generar_diagnostico_fechas_base_general() -> tuple[Path, dict[str, object]]:
    diagnostico_path = AUDITORIAS_DIR / f"diagnostico_fechas_base_general_{timestamp()}.csv"
    registros = leer_tsv_dicts(BASE_GENERAL)
    campos = [
        "timestamp",
        "fila",
        "tipo_documento",
        "num_documento",
        "dv",
        "campo_fecha",
        "valor",
        "clasificacion",
        "mensaje",
        "fuente_regla",
    ]
    conteos: Counter[str] = Counter()
    problemas_por_columna: Counter[str] = Counter()
    muestras: list[dict[str, str]] = []
    with diagnostico_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=campos)
        writer.writeheader()
        for fila, registro in enumerate(registros, start=2):
            for campo in sorted(COLUMNAS_FECHA):
                valor = registro.get(campo, "")
                clasificacion, _fecha, mensaje = clasificar_fecha_diagnostico(valor)
                conteos[clasificacion] += 1
                if clasificacion != "FECHA_OK":
                    problemas_por_columna[campo] += 1
                    if len(muestras) < 20:
                        muestras.append(
                            {
                                "fila": str(fila),
                                "documento": "|".join(
                                    [
                                        registro.get("TIPO_DOCUMENTO", ""),
                                        registro.get("NUM_DOCUMENTO", ""),
                                        registro.get("DV", ""),
                                    ]
                                ),
                                "campo": campo,
                                "valor": valor,
                                "clasificacion": clasificacion,
                                "accion_recomendada": accion_recomendada_fecha(clasificacion),
                            }
                        )
                writer.writerow(
                    {
                        "timestamp": datetime.now().isoformat(timespec="seconds"),
                        "fila": fila,
                        "tipo_documento": registro.get("TIPO_DOCUMENTO", ""),
                        "num_documento": registro.get("NUM_DOCUMENTO", ""),
                        "dv": registro.get("DV", ""),
                        "campo_fecha": campo,
                        "valor": valor,
                        "clasificacion": clasificacion,
                        "mensaje": mensaje,
                        "fuente_regla": FUENTE_PREVENTIVA,
                    }
                )
    resumen = {
        "diagnostico_fechas": rel(diagnostico_path),
        "total_filas_base_general": len(registros),
        "total_campos_fecha_revisados": len(registros) * len(COLUMNAS_FECHA),
        "conteos_fecha": dict(conteos),
        "problemas_por_columna": dict(problemas_por_columna),
        "muestras_problemas_fecha": muestras,
    }
    return diagnostico_path, resumen


def accion_recomendada_fecha(clasificacion: str) -> str:
    if clasificacion == "FECHA_NO_DISPONIBLE":
        return "Confirmar dato fuente o respaldar que no aplica; no inventar fecha."
    if clasificacion == "FECHA_FUTURA":
        return "Revisar contra documento fuente; confirmar si corresponde o corregir con respaldo."
    if clasificacion == "FECHA_AMBIGUA":
        return "Confirmar siglo/formato antes de normalizar."
    if clasificacion == "FECHA_INVALIDA":
        return "Corregir formato o valor usando respaldo documental."
    if clasificacion == "FECHA_VACIA":
        return "Confirmar si el campo puede quedar vacio."
    return "Sin accion."


def escribir_reporte(resumen: dict[str, object]) -> None:
    riesgos = []
    if int(resumen.get("total_posibles_duplicados_documento", 0)):
        riesgos.append("Existen documentos repetidos; no fueron eliminados automaticamente.")
    if int(resumen.get("total_valores_nd", 0)):
        riesgos.append("Existen valores #N/D que requieren confirmacion.")
    if int(resumen.get("total_valores_excel_na", 0)):
        riesgos.append("Existen valores #N/A preservados como datos fuente no disponibles.")
    if int(resumen.get("total_fechas_ambiguas", 0)) or int(resumen.get("total_fechas_futuras", 0)):
        riesgos.append("Existen fechas ambiguas o futuras que requieren revision.")
    if int(resumen.get("total_horas_coma_decimal", 0)):
        riesgos.append("Existen horas con coma decimal preservadas.")
    if not riesgos:
        riesgos.append("No se detectaron riesgos criticos en la incorporacion preliminar.")
    contenido = f"""# Reporte Fase 3 - Base Preliminar Personal Academico 2026

## Objetivo
Incorporar el Excel preliminar editable como RAW TSV, normalizarlo contra la estructura oficial de Personal Academico en la Institucion y alimentar una base general acumulable sin generar archivo final PES_READY.

## Resumen
- Fecha/hora: {resumen['fecha_hora']}
- Ruta Excel original: {EXCEL_ORIGINAL}
- Ruta copia Excel controlada: {rel(EXCEL_COPIA)}
- Hash SHA256 original: {resumen['hash_original']}
- Hash SHA256 copia: {resumen['hash_copia']}
- Hoja seleccionada: {resumen['hoja_seleccionada']}
- Total hojas detectadas: {resumen['total_hojas_detectadas']}
- Total filas RAW: {resumen['total_filas_raw']}
- Total filas utiles: {resumen['total_filas_utiles']}
- Total encabezados repetidos internos: {resumen['total_encabezados_repetidos_internos']}
- Total columnas RAW: {resumen['total_columnas_raw']}
- Total columnas oficiales: {len(COLUMNAS_OFICIALES)}
- Total filas normalizadas: {resumen['total_filas_normalizadas']}
- Total filas base general: {resumen['total_filas_base_general']}
- Total filas incorporadas: {resumen['total_filas_incorporadas']}
- Total filas actualizadas: {resumen['total_filas_actualizadas']}
- Total posibles duplicados por documento: {resumen['total_posibles_duplicados_documento']}
- Total documentos con multiples funciones/programas: {resumen['total_documentos_multiples_funciones']}
- Total valores #N/D: {resumen['total_valores_nd']}
- Total valores #N/A: {resumen['total_valores_excel_na']}
- Total valores N/A o NA: {resumen['total_valores_na']}
- Total valores NO APLICA: {resumen['total_valores_no_aplica']}
- Total campos fecha no disponibles: {resumen['total_campos_fecha_no_disponibles']}
- Total fechas ambiguas: {resumen['total_fechas_ambiguas']}
- Total fechas invalidas reales: {resumen['total_fechas_invalidas']}
- Total fechas futuras: {resumen['total_fechas_futuras']}
- Total horas con coma decimal: {resumen['total_horas_coma_decimal']}
- Total DV minuscula normalizada: {resumen['total_dv_minuscula_normalizada']}
- Total campos obligatorios preventivos vacios: {resumen['total_campos_obligatorios_preventivos_vacios']}

## Rutas Generadas
- RAW TSV: {rel(RAW_TSV)}
- Base normalizada: {rel(BASE_NORMALIZADA)}
- Base general: {rel(BASE_GENERAL)}
- Auditoria: {resumen['auditoria']}

## Riesgos Detectados
{chr(10).join(f'- {riesgo}' for riesgo in riesgos)}

## Proximos Pasos
- Revisar valores #N/D, fechas ambiguas/futuras y documentos repetidos.
- Confirmar criterios pendientes antes de exportacion final.
- Ejecutar exportacion final solo con auditoria aprobada y autorizacion explicita.

NO se genero archivo final PES_READY.

La base general sigue siendo preliminar y requiere revision antes de exportacion final.
"""
    REPORTE_MD.write_text(contenido, encoding="utf-8")


def escribir_reporte_3b(resumen: dict[str, object]) -> None:
    conteos = resumen.get("conteos_fecha", {})
    problemas = resumen.get("problemas_por_columna", {})
    muestras = resumen.get("muestras_problemas_fecha", [])
    problemas_txt = "\n".join(
        f"- {columna}: {total}" for columna, total in sorted(problemas.items(), key=lambda x: x[1], reverse=True)
    ) or "- No se detectaron problemas de fecha."
    muestras_txt = "\n".join(
        (
            f"- Fila {m['fila']} | Documento {m['documento']} | {m['campo']} | "
            f"`{m['valor']}` | {m['clasificacion']} | {m['accion_recomendada']}"
        )
        for m in muestras
    ) or "- Sin registros problematicos."
    contenido = f"""# Reporte Fase 3B - Depuracion de Fechas y Valores No Disponibles

## Objetivo
Distinguir tecnicamente fechas invalidas reales, fechas futuras, fechas ambiguas, campos de fecha vacios y valores no disponibles preservados como `#N/A`, `#N/D`, `N/A`, `NA` o `NO APLICA`.

## Resumen
- Fecha/hora: {resumen['fecha_hora']}
- Archivo revisado: {rel(BASE_GENERAL)}
- Cantidad de filas: {resumen['total_filas_base_general']}
- Campos de fecha revisados: {', '.join(sorted(COLUMNAS_FECHA))}
- Total campos de fecha revisados: {resumen['total_campos_fecha_revisados']}
- Total FECHA_OK: {conteos.get('FECHA_OK', 0)}
- Total FECHA_VACIA: {conteos.get('FECHA_VACIA', 0)}
- Total FECHA_NO_DISPONIBLE: {conteos.get('FECHA_NO_DISPONIBLE', 0)}
- Total FECHA_AMBIGUA: {conteos.get('FECHA_AMBIGUA', 0)}
- Total FECHA_INVALIDA: {conteos.get('FECHA_INVALIDA', 0)}
- Total FECHA_FUTURA: {conteos.get('FECHA_FUTURA', 0)}
- Total #N/A: {resumen.get('total_valores_excel_na', 0)}
- Total #N/D: {resumen.get('total_valores_nd', 0)}
- Total N/A o NA: {resumen.get('total_valores_na', 0)}
- Total NO APLICA: {resumen.get('total_valores_no_aplica', 0)}
- Diagnostico CSV: {resumen['diagnostico_fechas']}

## Columnas Con Mas Problemas
{problemas_txt}

## Muestra de Registros Problematicos
{muestras_txt}

## Criterio Aplicado
- `#N/A` no fue corregido.
- `#N/A` se preservo como dato fuente no disponible.
- Los valores no disponibles se clasifican como `FECHA_NO_DISPONIBLE` con fuente `VALIDACION_TECNICA_PREVENTIVA`.
- La base sigue preliminar.
- No se genero PES_READY.
"""
    REPORTE_3B_MD.write_text(contenido, encoding="utf-8")


def escribir_bitacora(resumen: dict[str, object]) -> None:
    contenido = f"""# Bitacora Fase 3 - Base Preliminar Personal Academico 2026

- Fecha/hora: {resumen['fecha_hora']}
- Rama git: {resumen['rama_git']}

## Comandos Ejecutados
- `git branch --show-current`
- `git switch -c feature/personal-academico-base-preliminar-2026`
- `mkdir -p ...`
- `cp /Users/alexi/Desktop/Libro1.xlsx ...`
- `shasum -a 256 ...`
- `python3 personal_academico_2026/scripts/incorporar_base_preliminar_personal_academico_2026.py`

## Archivos Creados
- {rel(EXCEL_COPIA)}
- {rel(RAW_TSV)}
- {rel(MAPEO_TSV)}
- {rel(BASE_NORMALIZADA)}
- {rel(BASE_GENERAL)}
- {resumen['auditoria']}
- {rel(REPORTE_MD)}
- {rel(BITACORA_MD)}
- {rel(RESUMEN_JSON)}

## Archivos Modificados
- {rel(BASE_GENERAL)}
- {rel(REPORTE_MD)}
- {rel(BITACORA_MD)}

## Archivos No Tocados
- No se tocaron CNED, matricula, scripts globales, Makefile global ni otros modulos.

## Resultado de Incorporacion
- Hoja seleccionada: {resumen['hoja_seleccionada']}
- Filas normalizadas: {resumen['total_filas_normalizadas']}
- Filas incorporadas: {resumen['total_filas_incorporadas']}
- Filas actualizadas: {resumen['total_filas_actualizadas']}

## Resultado de Validacion
- Estado incorporador: OK sin errores criticos.

## Advertencias
- Valores #N/D: {resumen['total_valores_nd']}
- Fechas ambiguas: {resumen['total_fechas_ambiguas']}
- Fechas futuras: {resumen['total_fechas_futuras']}
- Horas con coma decimal: {resumen['total_horas_coma_decimal']}
- Posibles duplicados por documento: {resumen['total_posibles_duplicados_documento']}
- Documentos con multiples funciones/programas: {resumen['total_documentos_multiples_funciones']}

## Errores
- No se registraron errores criticos.

## Estado Final
- Base general preliminar creada o actualizada.
- NO se genero archivo final PES_READY.

## Pendientes Antes de Exportar
- Depurar advertencias y confirmar reglas pendientes.
- Aprobar auditoria antes de generar archivo final PES.
"""
    BITACORA_MD.write_text(contenido, encoding="utf-8")


def escribir_bitacora_3b(resumen: dict[str, object], resultados_validacion: dict[str, str] | None = None) -> None:
    resultados_validacion = resultados_validacion or {}
    contenido = f"""# Bitacora Fase 3B - Depuracion de Fechas y NA

- Fecha/hora: {resumen['fecha_hora']}
- Rama: {resumen['rama_git']}

## Archivos Modificados
- {rel(BASE_NORMALIZADA)}
- {rel(BASE_GENERAL)}
- {rel(REPORTE_MD)}
- {rel(BITACORA_MD)}
- {rel(REPORTE_3B_MD)}
- {rel(BITACORA_3B_MD)}
- {rel(RESUMEN_JSON)}

## Archivos Creados
- {resumen['diagnostico_fechas']}
- {resumen['auditoria']}

## Comandos Ejecutados
- `git branch --show-current`
- `git status --short`
- `python3 -m py_compile personal_academico_2026/scripts/*.py`
- `python3 personal_academico_2026/scripts/incorporar_base_preliminar_personal_academico_2026.py`
- `python3 personal_academico_2026/scripts/validar_archivo_personal_academico_2026.py --tipo en_institucion --input personal_academico_2026/tests/fixtures/en_institucion_valido_minimo.csv`
- `python3 personal_academico_2026/scripts/validar_archivo_personal_academico_2026.py --tipo fuera_institucion --input personal_academico_2026/tests/fixtures/fuera_institucion_valido_minimo.csv`
- `python3 personal_academico_2026/scripts/validar_archivo_personal_academico_2026.py --tipo en_institucion --input personal_academico_2026/data/base_general/base_general_personal_academico_en_institucion.tsv --delimitador-entrada tab`
- `python3 personal_academico_2026/scripts/exportar_personal_academico_2026.py --tipo en_institucion --input personal_academico_2026/data/base_general/base_general_personal_academico_en_institucion.tsv --output personal_academico_2026/resultados/TEST_NO_DEBE_CREARSE_PES_READY.csv`

## Resultado Antes
- Base general TSV era rechazada por `fecha_formato` asociado principalmente a `#N/A`.

## Resultado Despues
- `#N/A`, `#N/D`, `N/A`, `NA` y `NO APLICA` se clasifican como no disponibles, no como formato invalido.
- Campos de fecha revisados: {resumen['total_campos_fecha_revisados']}
- FECHA_NO_DISPONIBLE: {resumen['conteos_fecha'].get('FECHA_NO_DISPONIBLE', 0)}
- FECHA_INVALIDA: {resumen['conteos_fecha'].get('FECHA_INVALIDA', 0)}
- FECHA_FUTURA: {resumen['conteos_fecha'].get('FECHA_FUTURA', 0)}

## Validaciones
- Fixture En Institucion: {resultados_validacion.get('fixture_en_institucion', 'pendiente')}
- Fixture Fuera Institucion: {resultados_validacion.get('fixture_fuera_institucion', 'pendiente')}
- Base general TSV: {resultados_validacion.get('base_general_tsv', 'pendiente')}
- Exportador bloqueado: {resultados_validacion.get('exportador_bloqueado', 'pendiente')}

## Pendientes
- Corregir o respaldar valores fuente no disponibles.
- Revisar fecha futura detectada.
- Mantener base como preliminar hasta depuracion y aprobacion.
- No generar PES_READY sin autorizacion y auditoria aprobada.
"""
    BITACORA_3B_MD.write_text(contenido, encoding="utf-8")


def git_branch() -> str:
    import subprocess

    proc = subprocess.run(
        ["git", "branch", "--show-current"],
        cwd=REPO_DIR,
        text=True,
        capture_output=True,
        check=False,
    )
    return proc.stdout.strip() or "desconocida"


def main() -> int:
    asegurar_directorios()
    ts = timestamp()
    auditoria_path = AUDITORIAS_DIR / f"auditoria_incorporacion_base_preliminar_{ts}.csv"
    auditor = Auditor(auditoria_path, EXCEL_ORIGINAL)
    try:
        hash_original, hash_copia, backup_excel = copiar_excel_controlado(auditor)
        wb = load_workbook(EXCEL_COPIA, read_only=True, data_only=True)
        hojas = auditar_hojas(wb, auditor)
        hoja_nombre = seleccionar_hoja(hojas, auditor)
        ws = wb[hoja_nombre]
        fila_encabezado, encabezado, repetidos = detectar_encabezado(ws, auditor)
        filas_raw = generar_raw(ws, fila_encabezado, encabezado, auditor)
        metricas_raw = validar_raw(encabezado, filas_raw, repetidos, auditor, hoja_nombre)
        asignaciones = escribir_catalogo_mapeo(encabezado, auditor, hoja_nombre)
        normalizadas, metricas_norm = normalizar_base(
            encabezado, filas_raw, asignaciones, auditor, hoja_nombre
        )
        metricas_general = escribir_base_general(normalizadas, auditor)
        diagnostico_path, metricas_diagnostico = generar_diagnostico_fechas_base_general()
        resumen = {
            "fecha_hora": datetime.now().isoformat(timespec="seconds"),
            "rama_git": git_branch(),
            "hash_original": hash_original,
            "hash_copia": hash_copia,
            "backup_excel": backup_excel,
            "total_hojas_detectadas": len(hojas),
            "hojas": hojas,
            "hoja_seleccionada": hoja_nombre,
            "auditoria": rel(auditoria_path),
            "raw_tsv": rel(RAW_TSV),
            "base_normalizada": rel(BASE_NORMALIZADA),
            "base_general": rel(BASE_GENERAL),
            "catalogo_mapeo": rel(MAPEO_TSV),
            "reporte": rel(REPORTE_MD),
            "bitacora": rel(BITACORA_MD),
            "reporte_3b": rel(REPORTE_3B_MD),
            "bitacora_3b": rel(BITACORA_3B_MD),
            "diagnostico_fechas": rel(diagnostico_path),
            "total_filas_normalizadas": len(normalizadas),
            **{k: 0 for k in [
                "total_valores_excel_na",
                "total_valores_nd",
                "total_valores_na",
                "total_valores_no_aplica",
                "total_campos_fecha_no_disponibles",
                "total_fechas_ambiguas",
                "total_fechas_futuras",
                "total_fechas_invalidas",
                "total_horas_coma_decimal",
                "total_dv_minuscula_normalizada",
                "total_campos_obligatorios_preventivos_vacios",
                "total_posibles_duplicados_documento",
                "total_documentos_multiples_funciones",
                "total_encabezados_repetidos_internos",
                "total_columnas_raw",
                "total_filas_raw",
                "total_filas_utiles",
                "total_filas_base_general",
                "total_filas_incorporadas",
                "total_filas_actualizadas",
            ]},
        }
        resumen.update(metricas_raw)
        resumen.update(metricas_norm)
        resumen.update(metricas_general)
        resumen.update(metricas_diagnostico)
        escribir_reporte(resumen)
        escribir_bitacora(resumen)
        escribir_reporte_3b(resumen)
        escribir_bitacora_3b(resumen)
        RESUMEN_JSON.write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")
        auditor.escribir()
        print("Incorporacion base preliminar Personal Academico 2026")
        print(f"estado: OK")
        print(f"hoja seleccionada: {hoja_nombre}")
        print(f"filas normalizadas: {len(normalizadas)}")
        print(f"base general: {rel(BASE_GENERAL)}")
        print(f"auditoria: {rel(auditoria_path)}")
        print("NO se genero archivo final PES_READY.")
        return 0
    except Exception as exc:
        auditor.add(
            "SIN_CAMBIOS",
            "ERROR",
            f"Incorporacion detenida por error critico: {exc}",
            FUENTE_PREVENTIVA,
            accion="DETENER",
        )
        auditor.escribir()
        print(f"ERROR: {exc}", file=sys.stderr)
        print(f"auditoria: {rel(auditoria_path)}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
