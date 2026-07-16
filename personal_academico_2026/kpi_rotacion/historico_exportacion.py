"""Exportacion y manifiestos del paquete historico de rotacion."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from .config import module_version, personal_project_root
from .historico_consolidacion import HistoricalTables
from .lectura import sha256_file


DESAGGREGATION_FILES = {
    "sexo": "13_rotacion_por_sexo.csv",
    "formacion": "14_rotacion_por_formacion.csv",
    "tramo_horas": "15_rotacion_por_tramo_horas.csv",
    "cargo": "16_rotacion_por_cargo.csv",
    "jerarquia": "17_rotacion_por_jerarquia.csv",
    "programa": "18_rotacion_por_programa.csv",
    "adscripcion": "19_rotacion_por_adscripcion.csv",
}


def timestamped_audit_dir(base_dir: Path | None = None, timestamp: datetime | None = None) -> Path:
    """Crea una carpeta de auditoria historica con marca de tiempo."""

    root = base_dir or personal_project_root() / "auditorias"
    value = timestamp or datetime.now()
    target = root / f"kpi_rotacion_historico_{value.strftime('%Y%m%d_%H%M%S')}"
    target.mkdir(parents=True, exist_ok=False)
    return target


def json_default(value: Any) -> Any:
    """Serializa objetos de pandas, pathlib y datetime para JSON."""

    if isinstance(value, Path):
        return str(value)
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    return str(value)


def safe_records(frame: pd.DataFrame, columns: list[str] | None = None) -> list[dict[str, object]]:
    """Convierte DataFrame a registros JSON, opcionalmente con columnas sanitizadas."""

    if frame.empty:
        return []
    selected = frame[columns].copy() if columns else frame.copy()
    return selected.where(pd.notna(selected), "").to_dict(orient="records")


def write_csv(path: Path, frame: pd.DataFrame) -> Path:
    """Escribe un CSV con separador institucional estable."""

    frame.to_csv(path, index=False, sep=";", encoding="utf-8-sig")
    return path


def format_workbook(path: Path) -> None:
    """Aplica formato minimo de auditoria al libro Excel generado."""

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
    workbook.save(path)


def executive_summary_frame(tables: HistoricalTables) -> pd.DataFrame:
    """Construye resumen ejecutivo sin datos personales."""

    rows = [
        {"INDICADOR": "VERSION_MODULO", "VALOR": module_version()},
        {"INDICADOR": "ANIOS", "VALOR": ",".join(map(str, tables.resumen.get("anios", [])))},
        {"INDICADOR": "PROMEDIO_SIMPLE_TASAS_ANUALES", "VALOR": tables.resumen.get("promedio_simple_tasas_anuales")},
        {"INDICADOR": "TASA_PONDERADA_HISTORICA", "VALOR": tables.resumen.get("tasa_ponderada_historica")},
        {"INDICADOR": "TOTAL_REINGRESOS", "VALOR": tables.resumen.get("total_reingresos")},
        {"INDICADOR": "FUENTES_INVENTARIADAS", "VALOR": len(tables.inventory)},
        {"INDICADOR": "FUENTES_ANUALES_SELECCIONADAS", "VALOR": len(tables.source_decisions)},
        {"INDICADOR": "PARES_ANUALES_EVALUADOS", "VALOR": len(tables.kpi)},
    ]
    return pd.DataFrame(rows)


def dictionary_frame() -> pd.DataFrame:
    """Documenta significado de salidas y campos principales."""

    rows = [
        ("TASA_ROTACION_SIES", "ROTAN / DOTACION_BASE * 100. Indicador oficial anual por par t,t+1."),
        ("NUEVOS_INGRESOS", "Universo t+1 menos universo t. KPI institucional; no participa en la tasa oficial."),
        ("PROMEDIO_SIMPLE_TASAS_ANUALES", "Media aritmetica de tasas anuales; no es tasa oficial consolidada."),
        ("TASA_PONDERADA_HISTORICA", "SUMA_ROTAN / SUMA_DOTACION_BASE * 100; resumen multianual institucional."),
        ("ESTADO_APTITUD", "Estado de publicabilidad/comparabilidad asignado de forma conservadora."),
        ("CLAVE_DOCUMENTAL", "TIPO_DOCUMENTO|NUM_DOCUMENTO_NORMALIZADO|DV_NORMALIZADO; identificador tecnico local."),
    ]
    return pd.DataFrame(rows, columns=["CAMPO", "DEFINICION"])


def export_excel(path: Path, tables: HistoricalTables, audit_frame: pd.DataFrame) -> Path:
    """Genera el Excel ejecutivo historico con hojas requeridas."""

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        executive_summary_frame(tables).to_excel(writer, sheet_name="00_RESUMEN_EJECUTIVO", index=False)
        tables.source_decisions.to_excel(writer, sheet_name="01_FUENTES_POR_ANIO", index=False)
        tables.kpi.to_excel(writer, sheet_name="02_KPI_HISTORICO", index=False)
        tables.detail.to_excel(writer, sheet_name="03_DETALLE_PARES_ANUALES", index=False)
        tables.panel.to_excel(writer, sheet_name="04_PANEL_PRESENCIA", index=False)
        tables.rotan.to_excel(writer, sheet_name="05_ROTAN", index=False)
        tables.permanecen.to_excel(writer, sheet_name="06_PERMANECEN", index=False)
        tables.nuevos.to_excel(writer, sheet_name="07_NUEVOS_INGRESOS", index=False)
        tables.reingresos.to_excel(writer, sheet_name="08_REINGRESOS", index=False)
        tables.cambios.to_excel(writer, sheet_name="09_CAMBIOS_DOCENTES", index=False)
        tables.desagregaciones.get("sexo", pd.DataFrame()).to_excel(writer, sheet_name="10_SEXO", index=False)
        tables.desagregaciones.get("formacion", pd.DataFrame()).to_excel(writer, sheet_name="11_FORMACION", index=False)
        tables.desagregaciones.get("tramo_horas", pd.DataFrame()).to_excel(writer, sheet_name="12_TRAMOS_HORAS", index=False)
        tables.desagregaciones.get("cargo", pd.DataFrame()).to_excel(writer, sheet_name="13_CARGO", index=False)
        tables.desagregaciones.get("jerarquia", pd.DataFrame()).to_excel(writer, sheet_name="14_JERARQUIA", index=False)
        tables.desagregaciones.get("programa", pd.DataFrame()).to_excel(writer, sheet_name="15_PROGRAMA", index=False)
        tables.desagregaciones.get("adscripcion", pd.DataFrame()).to_excel(writer, sheet_name="16_ADSCRIPCION", index=False)
        tables.identidad.to_excel(writer, sheet_name="17_VALIDACION_RUT", index=False)
        tables.duplicados.to_excel(writer, sheet_name="18_DUPLICADOS", index=False)
        tables.comparability.to_excel(writer, sheet_name="19_COMPARABILIDAD", index=False)
        audit_frame.to_excel(writer, sheet_name="20_AUDITORIA", index=False)
        dictionary_frame().to_excel(writer, sheet_name="21_DICCIONARIO", index=False)
    format_workbook(path)
    return path


def export_candidate_workbook(path: Path, tables: HistoricalTables) -> Path:
    """Genera el libro de comparacion de candidatos por anio."""

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        tables.source_decisions.to_excel(writer, sheet_name="00_DECISIONES", index=False)
        tables.candidate_comparison.to_excel(writer, sheet_name="01_CANDIDATOS", index=False)
        tables.inventory.to_excel(writer, sheet_name="02_INVENTARIO", index=False)
    format_workbook(path)
    return path


def manifest_sources(tables: HistoricalTables) -> dict[str, object]:
    """Construye manifiesto de fuentes sin datos personales."""

    inventory_columns = [
        "fuente_id",
        "relative_path",
        "extension",
        "size_bytes",
        "modified_at",
        "sha256",
        "year_from_path",
        "year_from_name",
        "year_from_content",
        "selected_sheet",
        "column_count",
        "physical_rows",
        "data_rows",
        "preliminary_classification",
        "evidence_pes",
        "evidence_finalized",
        "evidence_error",
        "observations",
    ]
    decision_columns = [
        "anio",
        "fuente_id",
        "hash",
        "personas_vigentes",
        "evidencia_principal",
        "evidencia_secundaria",
        "nivel_confianza",
        "estado",
        "aptitud_kpi",
        "observaciones",
    ]
    return {
        "version_modulo": module_version(),
        "inventario": safe_records(tables.inventory, [c for c in inventory_columns if c in tables.inventory.columns]),
        "decision_fuente_por_anio": safe_records(tables.source_decisions, [c for c in decision_columns if c in tables.source_decisions.columns]),
    }


def manifest_audit(
    output_dir: Path,
    tables: HistoricalTables,
    root: Path,
    desde_anio: int | None,
    hasta_anio: int | None,
    modo: str,
    commands: list[str],
    started_at: datetime,
    elapsed_seconds: float,
    artifacts: dict[str, Path],
) -> dict[str, object]:
    """Construye manifiesto tecnico de auditoria sin datos personales."""

    return {
        "fecha": datetime.now().date().isoformat(),
        "hora": datetime.now().time().replace(microsecond=0).isoformat(),
        "version_modulo": module_version(),
        "output_dir": str(output_dir),
        "historico_root": str(root),
        "desde_anio": desde_anio,
        "hasta_anio": hasta_anio,
        "modo": modo,
        "started_at": started_at.replace(microsecond=0).isoformat(),
        "elapsed_seconds": round(elapsed_seconds, 4),
        "commands": commands,
        "filas_longitudinales": len(tables.longitudinal),
        "personas_panel": len(tables.panel),
        "fuentes_inventariadas": len(tables.inventory),
        "fuentes_seleccionadas": len(tables.source_decisions),
        "pares_evaluados": len(tables.kpi),
        "duplicados_detectados": len(tables.duplicados),
        "conflictos_documentales": len(tables.conflictos),
        "artefactos": {
            key: {"path": str(path), "sha256": sha256_file(path)}
            for key, path in artifacts.items()
            if path.exists()
        },
    }


def json_summary(tables: HistoricalTables) -> dict[str, object]:
    """Construye JSON ejecutivo de KPI sin datos personales."""

    return {
        "version_modulo": module_version(),
        "formula_oficial": "ROTAN / DOTACION_BASE * 100, con ROTAN = presentes en t y ausentes en t+1.",
        "resumen_multianual": tables.resumen,
        "kpi_historico": safe_records(tables.kpi),
        "comparabilidad": safe_records(tables.comparability),
        "fuentes_por_anio": safe_records(
            tables.source_decisions,
            [
                "anio",
                "fuente_id",
                "hash",
                "personas_vigentes",
                "nivel_confianza",
                "estado",
                "aptitud_kpi",
                "observaciones",
            ],
        ),
    }


def markdown_report(tables: HistoricalTables, output_dir: Path) -> str:
    """Genera reporte Markdown ejecutivo sin datos personales."""

    lines = [
        "# KPI Rotacion Docente Historico",
        "",
        "## Formula Oficial SIES",
        "",
        "TASA_ROTACION_SIES(t,t+1) = academicos presentes en t y ausentes en t+1 / total de academicos presentes en t * 100.",
        "Los nuevos ingresos se informan como KPI institucional y no participan en numerador ni denominador.",
        "",
        "## Resumen",
        "",
        f"- Carpeta de auditoria: {output_dir}",
        f"- Fuentes inventariadas: {len(tables.inventory)}",
        f"- Fuentes anuales seleccionadas: {len(tables.source_decisions)}",
        f"- Pares anuales evaluados: {len(tables.kpi)}",
        f"- Promedio simple tasas anuales: {tables.resumen.get('promedio_simple_tasas_anuales')}",
        f"- Tasa ponderada historica: {tables.resumen.get('tasa_ponderada_historica')}",
        f"- Reingresos detectados: {tables.resumen.get('total_reingresos')}",
        "",
        "## Fuente Por Anio",
        "",
    ]
    for record in safe_records(tables.source_decisions):
        lines.append(
            f"- {record.get('anio')}: {record.get('fuente_id')} | estado={record.get('estado')} | aptitud={record.get('aptitud_kpi')} | personas_vigentes={record.get('personas_vigentes')}"
        )
    lines.extend(["", "## KPI Por Par", ""])
    for record in safe_records(tables.kpi):
        lines.append(
            f"- {record.get('ANIO_BASE')}-{record.get('ANIO_COMPARACION')}: base={record.get('DOTACION_BASE')} rotan={record.get('ROTAN')} tasa={record.get('TASA_ROTACION_SIES')} estado={record.get('ESTADO_APTITUD')}"
        )
    lines.extend(
        [
            "",
            "## Dictamen",
            "",
            "Resultado exploratorio gobernado. No publicar como KPI institucional definitivo hasta revisar evidencia PES finalizada y comparabilidad anual.",
        ]
    )
    return "\n".join(lines) + "\n"


def final_opinion(tables: HistoricalTables) -> str:
    """Construye dictamen final tecnico sin datos personales."""

    non_publicable = 0
    if not tables.kpi.empty and "ESTADO_APTITUD" in tables.kpi.columns:
        non_publicable = int(tables.kpi["ESTADO_APTITUD"].astype(str).ne("APTO_PARA_PUBLICACION").sum())
    lines = [
        "# Dictamen Final KPI Rotacion Historico",
        "",
        "El paquete historico fue generado en modo auditoria, sin integrar el modulo al flujo principal PES.",
        "",
        f"- Fuentes inventariadas: {len(tables.inventory)}",
        f"- Anios con fuente seleccionada: {len(tables.source_decisions)}",
        f"- Pares anuales calculados: {len(tables.kpi)}",
        f"- Pares no publicables o pendientes: {non_publicable}",
        "",
        "La tasa oficial SIES se conserva como indicador anual por par t,t+1. Los agregados multianuales son institucionales y estan etiquetados por separado.",
        "",
        "Recomendacion: revisar manualmente evidencia de carga PES finalizada por anio antes de publicar tasas institucionales o integrar al flujo principal.",
    ]
    return "\n".join(lines) + "\n"


def write_json(path: Path, payload: dict[str, object]) -> Path:
    """Escribe JSON con indentacion estable."""

    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=json_default) + "\n", encoding="utf-8")
    return path


def export_historical_package(
    tables: HistoricalTables,
    root: Path,
    desde_anio: int | None,
    hasta_anio: int | None,
    modo: str,
    commands: list[str],
    started_at: datetime,
    elapsed_seconds: float,
    output_base_dir: Path | None = None,
) -> dict[str, Path]:
    """Exporta todas las salidas historicas requeridas por la especificacion."""

    output_dir = timestamped_audit_dir(output_base_dir)
    artifacts: dict[str, Path] = {}

    csv_outputs = {
        "inventario": (output_dir / "01_inventario_fuentes_historicas.csv", tables.inventory),
        "decisiones": (output_dir / "02_decision_fuente_oficial_por_anio.csv", tables.source_decisions),
        "longitudinal": (output_dir / "04_personal_academico_historico_longitudinal.csv", tables.longitudinal),
        "panel": (output_dir / "05_panel_presencia_docente_por_anio.csv", tables.panel),
        "kpi_historico": (output_dir / "06_kpi_rotacion_docente_historico.csv", tables.kpi),
        "detalle": (output_dir / "07_detalle_rotacion_por_par_anual.csv", tables.detail),
        "permanecen": (output_dir / "08_permanecen_por_par_anual.csv", tables.permanecen),
        "rotan": (output_dir / "09_rotan_por_par_anual.csv", tables.rotan),
        "nuevos": (output_dir / "10_nuevos_ingresos_por_par_anual.csv", tables.nuevos),
        "reingresos": (output_dir / "11_reingresos_historicos.csv", tables.reingresos),
        "cambios": (output_dir / "12_cambios_docentes_entre_anios.csv", tables.cambios),
        "identidad": (output_dir / "20_validacion_identidad_por_rut.csv", tables.identidad),
        "duplicados": (output_dir / "21_duplicados_por_anio.csv", tables.duplicados),
        "conflictos": (output_dir / "22_conflictos_documentales.csv", tables.conflictos),
        "comparabilidad": (output_dir / "23_comparabilidad_pares_anuales.csv", tables.comparability),
    }
    for key, (path, frame) in csv_outputs.items():
        artifacts[key] = write_csv(path, frame)
    for key, filename in DESAGGREGATION_FILES.items():
        artifacts[key] = write_csv(output_dir / filename, tables.desagregaciones.get(key, pd.DataFrame()))

    artifacts["comparacion_candidatos"] = export_candidate_workbook(output_dir / "03_comparacion_candidatos_por_anio.xlsx", tables)
    audit_frame = pd.DataFrame(
        [
            {"CAMPO": "version_modulo", "VALOR": module_version()},
            {"CAMPO": "historico_root", "VALOR": str(root)},
            {"CAMPO": "desde_anio", "VALOR": desde_anio},
            {"CAMPO": "hasta_anio", "VALOR": hasta_anio},
            {"CAMPO": "modo", "VALOR": modo},
            {"CAMPO": "elapsed_seconds", "VALOR": round(elapsed_seconds, 4)},
        ]
    )
    artifacts["excel_historico"] = export_excel(output_dir / "KPI_ROTACION_DOCENTE_HISTORICO.xlsx", tables, audit_frame)
    artifacts["kpi_json"] = write_json(output_dir / "24_kpi_rotacion_historico.json", json_summary(tables))
    artifacts["kpi_md"] = output_dir / "25_kpi_rotacion_historico.md"
    artifacts["kpi_md"].write_text(markdown_report(tables, output_dir), encoding="utf-8")
    artifacts["manifest_fuentes"] = write_json(output_dir / "26_manifest_fuentes.json", manifest_sources(tables))

    artifacts["resultado_pruebas"] = output_dir / "28_resultado_pruebas.txt"
    artifacts["resultado_pruebas"].write_text(
        "Pendiente de actualizar despues de ejecutar compileall/pytest en el entorno gobernado.\n",
        encoding="utf-8",
    )
    artifacts["dictamen"] = output_dir / "29_dictamen_final.md"
    artifacts["dictamen"].write_text(final_opinion(tables), encoding="utf-8")

    audit_payload = manifest_audit(
        output_dir=output_dir,
        tables=tables,
        root=root,
        desde_anio=desde_anio,
        hasta_anio=hasta_anio,
        modo=modo,
        commands=commands,
        started_at=started_at,
        elapsed_seconds=elapsed_seconds,
        artifacts=artifacts,
    )
    artifacts["manifest_auditoria"] = write_json(output_dir / "27_manifest_auditoria.json", audit_payload)
    artifacts["output_dir"] = output_dir
    return artifacts
