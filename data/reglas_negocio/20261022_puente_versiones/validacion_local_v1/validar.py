"""Validación local, agregada y no destructiva. Ejecutar desde la raíz del repo."""
import csv,json,hashlib,re,sqlite3,subprocess
from pathlib import Path
from collections import defaultdict,Counter
from datetime import datetime,timezone
import openpyxl
ROOT=Path.cwd(); BASE=ROOT/'data/reglas_negocio/20261022_puente_versiones'; OUT=BASE/'validacion_local_v1'
CAT=ROOT/'data/catalogos_sies/2027/5910_20260907_6047_20261013_modalidad_v2/catalogo.json'
HIST=Path('/Users/alexi/Documents/GitHub/CARACTERIZACIÓN/gobernanza/sies/historico/2007_2026_h5_20261015/maestro.json')
PRO=Path('/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_12166.xlsx')
MU=ROOT/'outputs/base_retencion_matricula_unificada_2022_2026/20260901_142342/BASE_RETENCION_MATRICULA_UNIFICADA_INFORMADA_SIES_2022_2026_20260901_142342.xlsx'
def sha(p):
 h=hashlib.sha256()
 with open(p,'rb') as f:
  for b in iter(lambda:f.read(1048576),b''):h.update(b)
 return h.hexdigest()
def js(name,obj): (OUT/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf8')
def csvout(name,rows):
 if not rows:return
 with open(OUT/name,'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def readcsv(p):return list(csv.DictReader(open(p,encoding='utf-8-sig')))
def val(x):
 if x is None:return ''
 if isinstance(x,float) and x.is_integer():return str(int(x))
 return str(x).strip().upper()
def rut(a,b):
 a=val(a).replace('.','').replace(' ',''); b=val(b).replace(' ','')
 if not re.fullmatch(r'\d+',a) or not re.fullmatch(r'[0-9K]',b):return None
 return str(int(a))+'-'+b

def readxl(p,sheet=None):
 w=openpyxl.load_workbook(p,read_only=True,data_only=True);s=w[sheet] if sheet else w.active
 it=iter(s.values);header=[val(x) for x in next(it)]
 for row in it:yield dict(zip(header,row))
 w.close()
def git(p):
 return {k:subprocess.check_output(['git','-C',str(p),*args],text=True).strip() for k,args in {'rama':['branch','--show-current'],'HEAD':['rev-parse','HEAD'],'remotos':['remote','-v'],'status':['status','--short']}.items()}
repos=[ROOT,HIST.parents[4]]
initial_git={str(p):git(p) for p in repos};js('GIT_ANTES.json',initial_git)
files=[CAT,HIST,PRO,MU]+sorted(p for p in BASE.iterdir() if p.is_file())
# Preservar también los resultados locales existentes de retención, sin recalcularlos.
retfiles=sorted(p for p in (ROOT/'outputs/retencion_intermedia_t2_t3').rglob('*') if p.is_file() and p.suffix in {'.json','.csv'} and 'node_modules' not in p.parts)
ret_hash={str(p):sha(p) for p in retfiles}
inv=[]
for p in files:
 assert p.exists(),f'FUENTE_ESENCIAL_AUSENTE: {p}'
 r={'ruta':str(p),'bytes':p.stat().st_size,'sha256':sha(p),'clasificacion':'B'}
 if p.suffix=='.xlsx':
  w=openpyxl.load_workbook(p,read_only=True,data_only=True)
  r['hojas']=[{'nombre':s.title,'filas_incluye_encabezado':s.max_row,'columnas':[str(x) for x in next(s.values)]} for s in w];w.close()
 elif p.suffix=='.json':r['claves']=list(json.load(open(p)))
 elif p.suffix=='.csv':r['columnas']=list(readcsv(p)[0])
 inv.append(r)
js('FASE0_FUENTES.json',inv)
cat=json.load(open(CAT));hist=json.load(open(HIST));old=readcsv(BASE/'DICCIONARIO_VERSIONES_UMAS_SIES_PROPUESTO_20261022.csv');bridge=readcsv(BASE/'PUENTE_VERSIONES_PLAN_UMAS_SIES_20261022.csv')
dims={(r['codcarpr_interno'],r['sede'],r['jornada_catalogo']):r for r in old}
candidates=defaultdict(set)
for r in bridge:candidates[(r['codcarpr_interno'],r['sede'],r['jornada_catalogo'])].add(r['codigo_unico_sies_candidato'])
assert len(dims)==63
mi=defaultdict(set); my=defaultdict(set); mu_counts=Counter();mu_people=defaultdict(set);types=Counter();mu_bad=0
for r in readxl(MU,'BASE_RETENCION_MU'):
 types[val(r['TIPO_DOC'])]+=1
 code='I162S%sC%sJ%sV%s'%tuple(val(r[k]) for k in ['COD_SED','COD_CAR','JOR','VERSION']);y=val(r['ANIO_INFORMADO']);k=rut(r['N_DOC'],r['DV'])
 for axis,item in [('anio',y),('anio_codigo',y+'|'+code),('anio_carrera',y+'|'+val(r['COD_CAR'])),('anio_sede_jornada',y+'|'+val(r['COD_SED'])+'|'+val(r['JOR']))]:mu_counts[(axis,item)]+=1
 if val(r['VIG'])=='1':mu_counts[('vig1',y)]+=1
 if k and val(r['TIPO_DOC'])=='R':mi[k].add(code);my[(k,y)].add(code);mu_people[code].add(k)
 else:mu_bad+=1
ps=set();codcli_map=defaultdict(set); states=Counter();title=defaultdict(set);years=set();bad=0;raw=0
for r in readxl(PRO):
 raw+=1;k=rut(r['RUT'],r['DIG']);c=val(r['CODCARR']);p=val(r['PLAN_DE_ESTUDIO']);y=val(r['ANO']);t=val(r['PERIODO']);years.add(y);states[val(r['ESTADO_ACADEMICO'])]+=1
 if not k:bad+=1;continue
 ps.add((k,c,p,y,t));codcli_map[(val(r['CODCLI']),p)].add(k)
 if 'TITUL' in val(r['ESTADO_ACADEMICO']):title[y].add(k)
# Evidencia histórica sin adjudicar año; se conservan todos los códigos exactos observados.
dist={};dim_people={};agg=defaultdict(set);global_groups=defaultdict(set)
for k,c,p,y,t in ps:
 global_groups[(c,p,y,t)].add(k)
 for code in my.get((k,y),set()) or {''}:agg[(c,p,y,t,code)].add(k)
for d in dims:
 people={r[0] for r in ps if r[1]==d[0]};dim_people[d]=people
 dist[d]={code:len(people & mu_people[code]) for code in sorted(candidates[d]) if people & mu_people[code]}
csvout('FASE1_CRUCE_AGREGADO.csv',[{'carrera_interna':c,'plan':p,'ano':y,'periodo':t,'codigo_unico_sies':code,'personas':len(v),'coincidencias':len(v) if code else 0,'sin_match_mismo_ano':0 if code else len(v),'clasificacion':'C','ambito':'MISMO_ANIO; sin match no implica desercion'} for (c,p,y,t,code),v in sorted(agg.items())])
coverage={'PROMEDIOS_anios':sorted(years),'MU_anios':sorted(k[1] for k in mu_counts if k[0]=='anio'),'TIPO_DOC_MU':dict(types),'filas_PROMEDIOS':raw,'personas_plan_ano_periodo_carrera':len(ps),'filas_repetidas_en_grano':raw-bad-len(ps),'formato_no_utilizable_PROMEDIOS_filas':bad,'MU_no_RUT_o_formato_no_utilizable_filas':mu_bad,'CODCLI_plan_con_varios_RUT':sum(len(v)>1 for v in codcli_map.values()),'estados_academicos_filas':dict(states),'llaves':{'PROMEDIOS':'RUT/DIG; CODCLI+PLAN_DE_ESTUDIO solo relacion interna','MU':'TIPO_DOC=R,N_DOC,DV; ANIO_INFORMADO para cruce temporal'},'oferta_2027':len(cat['oferta']),'oferta_historica_anios':sorted({r['anio'] for r in hist['oferta_anual']}),'clasificacion':'B/C'}
js('COBERTURA_Y_LLAVES.json',coverage)
offer={r['codigo_unico']:r for r in cat['oferta']};history=defaultdict(list)
for r in hist['oferta_anual']:history[r['codigo_unico']].append(r)
allrows=[];review=[];nomatch=[];differences=[]
for d,prev in dims.items():
 counts=dist[d];total=sum(counts.values());unique=len(set().union(*(dim_people[d]&mu_people[x] for x in counts))) if counts else 0
 pred=max(counts,key=counts.get) if counts else '';ratio=counts.get(pred,0)/total if total else None
 oldcounts={a.rsplit(':',1)[0]:int(a.rsplit(':',1)[1]) for a in prev['distribucion_versiones'].split(';') if a}
 if counts!=oldcounts:differences.append({'dimension':list(d),'anterior':oldcounts,'reproducido':counts})
 plans=sorted({r[2] for r in ps if r[1]==d[0]});ys=sorted({r[3] for r in ps if r[1]==d[0]});h=history.get(pred,[])
 row={'carrera_interna':d[0],'sede':d[1],'jornada':d[2],'planes':';'.join(plans),'primer_ano_PROMEDIOS':min(ys) if ys else '', 'candidatos':';'.join(sorted(candidates[d])),'distribucion_versiones':';'.join(f'{k}:{v}' for k,v in counts.items()),'personas_coincidentes_unicas':unique,'incidencias_persona_codigo':total,'personas_multicodigo':sum(sum(k in mu_people[x] for x in counts)>1 for k in dim_people[d]),'codigo_predominante_observado':pred,'version_observada':pred.rsplit('V',1)[-1] if pred else '', 'proporcion_incidencias':ratio,'estado_anterior':prev['estado'],'estado':'PENDIENTE','clasificacion':'C/E','tipo_plan_SIES':offer.get(pred,{}).get('plan',''),'duracion_estudios_SIES':offer.get(pred,{}).get('duracion_estudios',''),'primer_ano_oferta':min(r['anio'] for r in h) if h else '', 'vigencia_historica':';'.join(f"{r['anio']}:{r['vigencia']}" for r in sorted(h,key=lambda r:r['anio'])),'tipo_duracion_plan_UMAS':'NO_ACREDITADOS','control_MU_personas_candidatos':len(set().union(*(mu_people[x] for x in candidates[d]))),'pendiente':'Acreditar tipo/duracion UMAS, temporalidad y conciliacion de universo; no promover por predominio','evidencia':'FASE1_CRUCE_AGREGADO.csv; FASE0_FUENTES.json'}
 allrows.append(row)
 if prev['estado']=='PENDIENTE_MULTIPLE_O_BAJA_COBERTURA':review.append(row)
 if prev['estado']=='PENDIENTE_SIN_MATCH':
  pe=dim_people[d];anymatch=pe&set(mi);samebody={k.split('-')[0] for k in pe}&{k.split('-')[0] for k in mi}; backup=[p for p in hist['planes_umas'] if d[0] in p.get('carreras_promedios',[]) or any(x.get('Codcarr')==d[0] for x in p.get('contextos_catalogo',[]))]
  nomatch.append({'carrera_interna':d[0],'sede':d[1],'jornada':d[2],'planes':';'.join(plans),'anios_PROMEDIOS':';'.join(ys),'anios_MU':'2022;2023;2024;2025;2026','personas_PROMEDIOS':len(pe),'personas_match_MU_cualquier_codigo':len(anymatch),'personas_match_candidatos':unique,'cuerpos_coincidentes_sin_DV':len(samebody),'planes_en_respaldo_gobernado':len(backup),'candidatos':row['candidatos'],'causa':'cobertura temporal' if ys and all(y<'2022' for y in ys) else 'ausencia de evidencia','observacion':'No hay PLAN_DE_ESTUDIO en MU; identidad coincidente en otro codigo no demuestra equivalencia. El respaldo registra actividad, no asignacion acreditada.','clasificacion':'B/C/E'})
csvout('TABLA_63_DIMENSIONES.csv',allrows);csvout('FASE2_REVISION_11.csv',review);csvout('FASE3_DIAGNOSTICO_13.csv',nomatch)
js('DIFERENCIAS_DERIVADO_PREVIO.json',differences)
print(json.dumps({'fase':'1-3','dimensiones':len(allrows),'revision':len(review),'sin_match':len(nomatch),'diferencias_previo':len(differences),'tipos_documento':dict(types),'cobertura':sorted(years)},ensure_ascii=False))
# Segunda lectura de fuentes, implementación relacional independiente en memoria.
db=sqlite3.connect(':memory:');db.executescript('CREATE TABLE p (id TEXT,c TEXT,plan TEXT,y TEXT,t TEXT);CREATE TABLE m (id TEXT,y TEXT,code TEXT,vig TEXT,car TEXT,sede TEXT,jor TEXT);CREATE TABLE cand (c TEXT,s TEXT,j TEXT,code TEXT);')
def norm2(a,b):
 a=str(a if a is not None else '').strip().upper();b=str(b if b is not None else '').strip().upper()
 if a.endswith('.0'):a=a[:-2]
 if b.endswith('.0'):b=b[:-2]
 a=a.replace('.','').replace(' ','');b=b.replace(' ','')
 return (str(int(a))+'-'+b) if a.isdecimal() and len(b)==1 and b in '0123456789K' else None
for r in readxl(PRO):
 ident=norm2(r['RUT'],r['DIG'])
 if ident:db.execute('INSERT INTO p VALUES (?,?,?,?,?)',(ident,val(r['CODCARR']),val(r['PLAN_DE_ESTUDIO']),val(r['ANO']),val(r['PERIODO'])))
for r in readxl(MU,'BASE_RETENCION_MU'):
 code='I162S'+val(r['COD_SED'])+'C'+val(r['COD_CAR'])+'J'+val(r['JOR'])+'V'+val(r['VERSION'])
 db.execute('INSERT INTO m VALUES (?,?,?,?,?,?,?)',(norm2(r['N_DOC'],r['DV']) if val(r['TIPO_DOC'])=='R' else None,val(r['ANIO_INFORMADO']),code,val(r['VIG']),val(r['COD_CAR']),val(r['COD_SED']),val(r['JOR'])))
for d,cs in candidates.items():db.executemany('INSERT INTO cand VALUES (?,?,?,?)',[(*d,c) for c in cs])
db.executescript('CREATE INDEX ip ON p(id,c); CREATE INDEX im ON m(id,code);')
sql_dist=defaultdict(dict)
for c,s,j,code,n in db.execute('SELECT cand.c,cand.s,cand.j,cand.code,COUNT(DISTINCT p.id) FROM cand JOIN p ON p.c=cand.c JOIN m ON m.id=p.id AND m.code=cand.code GROUP BY cand.c,cand.s,cand.j,cand.code'):sql_dist[(c,s,j)][code]=n
sqlagg={(c,p,y,t,code or ''):n for c,p,y,t,code,n in db.execute('SELECT p.c,p.plan,p.y,p.t,m.code,COUNT(DISTINCT p.id) FROM p LEFT JOIN m ON m.id=p.id AND m.y=p.y GROUP BY p.c,p.plan,p.y,p.t,m.code')}
second=Counter()
for axis,expr,where in [('anio','y',''),('vig1','y',"WHERE vig='1'"),('anio_codigo',"y||'|'||code",''),('anio_carrera',"y||'|'||car",''),('anio_sede_jornada',"y||'|'||sede||'|'||jor",'')]:
 for key,n in db.execute(f'SELECT {expr},COUNT(*) FROM m {where} GROUP BY {expr}'):second[(axis,key)]=n
controls=[{'indicador':a,'grupo':b,'antes':n,'despues':second[(a,b)],'brecha':second[(a,b)]-n,'clasificacion':'C'} for (a,b),n in sorted(mu_counts.items())]
csvout('FASE4_TOTALES_Y_BRECHAS.csv',controls)
csvout('TITULADOS_OBSERVADOS.csv',[{'ano':y,'personas_tituladas_observadas':len(v),'referencia_oficial':'PENDIENTE','clasificacion':'B/C; no cohorte de titulacion'} for y,v in sorted(title.items())])
ret_ok=all(p.exists() and sha(p)==h for p,h in ((Path(p),h) for p,h in ret_hash.items()))
js('PRESERVACION_RETENCION.json',{'clasificacion':'C/E','archivos_antes':ret_hash,'hashes_iguales_despues':ret_ok,'alcance':'No se recalcularon denominadores, retenidos ni tasas. Se preservaron bytes de resultados previos. Conciliacion semantica con puente PENDIENTE.'})
source_ok=all(sha(Path(r['ruta']))==r['sha256'] for r in inv)
qa={'clasificacion':'C/E','implementacion_1':'Python conjuntos, lectura openpyxl','implementacion_2':'Segunda lectura XLSX + SQLite en memoria, normalizador independiente','dimensiones':len(dims),'candidatos':sum(map(len,candidates.values())),'distribuciones_iguales':all(dist[d]==sql_dist[d] for d in dims),'agregados_iguales':{k:len(v) for k,v in agg.items()}==sqlagg,'totales_iguales':mu_counts==second,'hashes_fuentes_iguales':source_ok,'retencion_bytes_preservados':ret_ok,'diferencias_con_derivado_previo':len(differences),'asignaciones_nuevas':0,'pendientes_validacion':len(dims),'clasificacion_previa':dict(Counter(r['estado'] for r in old)),'duplicados_grano':raw-bad-len(ps),'sql_duplicados_grano':db.execute('SELECT COUNT(*) FROM p').fetchone()[0]-db.execute('SELECT COUNT(*) FROM (SELECT DISTINCT * FROM p)').fetchone()[0],'personas_sin_match_mismo_ano_en_grupos':sum(len(v) for k,v in agg.items() if not k[-1]),'sql_personas_sin_match_mismo_ano_en_grupos':sum(n for k,n in sqlagg.items() if not k[-1]),'limitacion':'QA computacional no acredita tipo de plan ni equivalencia oficial. Personas repetidas entre celdas no son aditivas.'}
qa['QA_COMPUTACIONAL']='PASS' if all(qa[k] for k in ['distribuciones_iguales','agregados_iguales','totales_iguales','hashes_fuentes_iguales','retencion_bytes_preservados']) and qa['duplicados_grano']==qa['sql_duplicados_grano'] else 'FAIL'
qa['PUERTA_PROMOCION']='BLOQUEADA';js('FASE5_QA_INDEPENDIENTE.json',qa)
print(json.dumps(qa,ensure_ascii=False))
js('GIT_DESPUES.json',{str(p):git(p) for p in repos})
js('MANIFIESTO_AUDITORIA.json',{'fecha':datetime.now(timezone.utc).isoformat(),'estado':'PENDIENTE_APROBACION_GOBERNANZA','fuentes':'FASE0_FUENTES.json','metodo':'validar.py; conjuntos Python y SQL independiente','alcance':'Auditoria de 63 dimensiones. No diccionario definitivo ni equivalencia oficial.','privacidad':'Cruces individuales en memoria local; salidas solo agregadas sin RUT, CODCLI, nombres o nacimiento. Sin servicios externos.','responsable_ejecucion':'Codex por encargo del usuario','responsable_aprobacion':'Gobernanza IPSS, pendiente de designacion','limitaciones':['Tipo y duracion UMAS no acreditados','MU no informa plan','Cruce historico no prueba equivalencia temporal','Retencion preservada por hash; conciliacion semantica no acreditada'],'clasificacion':'C/D/E','commit_push':False})
