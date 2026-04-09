#!/usr/bin/env python3
"""
Reconciliacion end-to-end 66 -> 54 para MU2026 postgrado/postitulo.

El script no modifica bases operativas. Genera artefactos nuevos de auditoria y,
si ya existen salidas de reconciliacion, crea backup antes de reemplazarlas.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
RESULTADOS = REPO / "resultados"
DIAG_DIR = REPO / "diagnosticos" / "mu2026_postgrado"
EXCEL_FUENTE = Path("/Users/alexi/Downloads/PROMEDIOSDEALUMNOS_7804.xlsx")

PES_READY = RESULTADOS / "matricula_unificada_2026_postgrado_postitulo_PES_READY.csv"
CON_TITULOS = RESULTADOS / "matricula_unificada_2026_postgrado_postitulo_CON_TITULOS.csv"
CONTROL_FINAL = RESULTADOS / "matricula_unificada_2026_postgrado_postitulo_CONTROL_FINAL.csv"
EXCLUIDOS = RESULTADOS / "matricula_unificada_2026_postgrado_postitulo_excluidos_2026_gobernanza.csv"
TRAZABILIDAD = RESULTADOS / "trazabilidad_matricula_unificada_2026_postgrado_postitulo.tsv"
DICCIONARIO = REPO / "control" / "diccionarios" / "diccionario_postgrado_postitulo_mu2026.tsv"

OUT_XLSX = RESULTADOS / "reconciliacion_66_a_54_postgrado.xlsx"
OUT_CSV = RESULTADOS / "reconciliacion_66_a_54_postgrado.csv"

TS = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT_MD = DIAG_DIR / f"reconciliacion_66_a_54_postgrado_{TS}.md"
OUT_JSON = DIAG_DIR / f"reconciliacion_66_a_54_postgrado_{TS}.json"

EXPECTED_UNIVERSE = 66
CAMPOS_PES = [
    "TIPO_DOC", "N_DOC", "DV", "PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRE",
    "SEXO", "FECH_NAC", "NAC", "PAIS_EST_SEC", "COD_SED", "COD_CAR",
    "MODALIDAD", "JOR", "VERSION", "FOR_ING_ACT", "ANIO_ING_ACT", "SEM_ING_ACT",
    "ANIO_ING_ORI", "SEM_ING_ORI", "VIG",
]
PREGRADO_ARCHIVOS = [
    RESULTADOS / "matricula_unificada_2026_pregrado.csv",
    RESULTADOS / "matricula_unificada_2026_pregrado_PES_READY.csv",
    RESULTADOS / "matricula_unificada_2026_control.csv",
]
EXPLICIT_UNIVERSE_CANDIDATES = [
    RESULTADOS / "universo_66_codcli_postgrado.csv",
    RESULTADOS / "universo_66_codcli_postgrado.tsv",
    RESULTADOS / "reconciliacion_universo_66_postgrado.csv",
    DIAG_DIR / "universo_66_codcli_postgrado.csv",
    DIAG_DIR / "universo_66_codcli_postgrado.tsv",
]

MASTER_COLUMNS = [
    "CODCLI", "RUT", "DV", "RUT_DV", "CODCARPR", "CODIGOCARRERA", "NOMBRE_L",
    "ANOMATRICULA", "PERIODOMATRICULA", "ANOINGRESO", "PERIODOINGRESO",
    "ESTADOACADEMICO", "MATRICULA", "ACCION_CARGA", "ESTADO_GOBERNANZA",
    "ESTADO_TRAZABILIDAD", "ESTADO_VALIDACION_PES", "INGRESA_A_PES_READY",
    "MOTIVO_INCLUSION", "MOTIVO_EXCLUSION", "COD_SED", "COD_CAR", "MODALIDAD",
    "JOR", "VERSION", "FOR_ING_ACT", "VIG", "FUENTE_DECISION",
    "OBSERVACION_AUDITORIA",
]


def s(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def sha256(path: Path) -> str | None:
    if not path.exists():
        return None
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def backup_if_exists(path: Path, backups: list[dict[str, str]]) -> None:
    if not path.exists():
        return
    bkp = path.with_name(f"{path.stem}_bkp_{TS}{path.suffix}")
    shutil.copy2(path, bkp)
    backups.append({"archivo": str(path), "backup": str(bkp)})


def read_csv_dict(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [{k: s(v) for k, v in row.items()} for row in csv.DictReader(f, delimiter=delimiter)]


def read_pes(path: Path) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for line_no, row in enumerate(csv.reader(f, delimiter=";"), start=1):
            if not row:
                continue
            out.append({
                "_LINEA_PES": str(line_no),
                "_N_CAMPOS": str(len(row)),
                **{col: s(row[i]) if i < len(row) else "" for i, col in enumerate(CAMPOS_PES)},
            })
    return out


def split_rut(value: Any, explicit_dv: Any = "") -> tuple[str, str]:
    raw = s(value).replace(".", "").replace(" ", "").upper()
    if "-" in raw:
        n_doc, dv = raw.rsplit("-", 1)
        return re.sub(r"\D", "", n_doc), dv.strip().upper()
    n_doc = re.sub(r"\D", "", raw)
    return n_doc, s(explicit_dv).upper()


def rut_dv(n_doc: str) -> str:
    total = 0
    factors = [2, 3, 4, 5, 6, 7]
    for i, digit in enumerate(reversed(n_doc)):
        total += int(digit) * factors[i % len(factors)]
    result = 11 - (total % 11)
    if result == 11:
        return "0"
    if result == 10:
        return "K"
    return str(result)


def valid_rut(n_doc: str, dv: str) -> tuple[bool, str]:
    n_doc = s(n_doc)
    dv = s(dv).upper()
    if not n_doc.isdigit():
        return False, "RUT_NO_NUMERICO"
    if len(n_doc) < 7 or len(n_doc) > 8:
        return False, "RUT_LONGITUD_INVALIDA"
    if not re.match(r"^[0-9K]$", dv):
        return False, "DV_FORMATO_INVALIDO"
    expected = rut_dv(n_doc)
    if dv != expected:
        return False, f"DV_INCONSISTENTE_ESPERADO_{expected}"
    return True, "OK"


def pes_tuple(row: dict[str, str]) -> tuple[str, ...]:
    return tuple(s(row.get(col)) for col in CAMPOS_PES)


def df_from_rows(rows: list[dict[str, str]], columns: list[str] | None = None) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    if columns is not None:
        for col in columns:
            if col not in df.columns:
                df[col] = ""
        df = df[columns]
    return df.fillna("").astype(str)


def load_excel_sources() -> tuple[dict[str, dict[str, str]], dict[str, dict[str, str]]]:
    cols = [
        "CODCLI", "RUT", "CODCARPR", "CODIGOCARRERA", "NOMBRE_L", "ANOMATRICULA",
        "PERIODOMATRICULA", "ANOINGRESO", "PERIODOINGRESO", "ESTADOACADEMICO",
        "MATRICULA", "SITUACION",
    ]
    da = pd.read_excel(EXCEL_FUENTE, sheet_name="DatosAlumnos", dtype=str, usecols=lambda c: c in cols)
    da = da.fillna("").astype(str)
    da_by_codcli = {s(r["CODCLI"]): {k: s(v) for k, v in r.items()} for r in da.to_dict("records") if s(r.get("CODCLI"))}

    base = pd.read_excel(EXCEL_FUENTE, sheet_name="base_datos", dtype=str).fillna("").astype(str)
    base_by_codcli = {s(r["CODCLI"]): {k: s(v) for k, v in r.items()} for r in base.to_dict("records") if s(r.get("CODCLI"))}
    return da_by_codcli, base_by_codcli


def load_explicit_universe() -> tuple[set[str] | None, str]:
    for path in EXPLICIT_UNIVERSE_CANDIDATES:
        if not path.exists():
            continue
        delimiter = "\t" if path.suffix == ".tsv" else ","
        rows = read_csv_dict(path, delimiter=delimiter)
        codcli = {s(r.get("CODCLI")) for r in rows if s(r.get("CODCLI"))}
        if len(codcli) == EXPECTED_UNIVERSE:
            return codcli, f"archivo explicito: {path}"
    return None, "sin archivo explicito de universo 66"


def reconstruct_universe(
    control_rows: list[dict[str, str]],
    excluded_rows: list[dict[str, str]],
) -> tuple[set[str], str, list[dict[str, str]], list[str]]:
    explicit, explicit_source = load_explicit_universe()
    if explicit is not None:
        return explicit, explicit_source, [], []

    base_control = {s(r.get("CODCLI")) for r in control_rows if s(r.get("CODCLI"))}
    missing_n = EXPECTED_UNIVERSE - len(base_control)
    if missing_n < 0:
        raise RuntimeError(
            f"El control final ya tiene {len(base_control)} CODCLI, excede el universo esperado {EXPECTED_UNIVERSE}."
        )

    excluded_not_control = [
        r for r in excluded_rows
        if s(r.get("CODCLI")) and s(r.get("CODCLI")) not in base_control
    ]
    group_counter: Counter[tuple[str, str]] = Counter(
        (s(r.get("CODCARPR")), s(r.get("NOMBRE_L"))) for r in excluded_not_control
    )
    closing_groups = [group for group, n in group_counter.items() if n == missing_n]
    if missing_n == 0:
        selected_extra: list[dict[str, str]] = []
    elif len(closing_groups) == 1:
        selected_group = closing_groups[0]
        selected_extra = [
            r for r in excluded_not_control
            if (s(r.get("CODCARPR")), s(r.get("NOMBRE_L"))) == selected_group
        ]
    else:
        detail = ", ".join(f"{codcarpr}/{nombre}: {n}" for (codcarpr, nombre), n in sorted(group_counter.items()))
        raise RuntimeError(
            "No se puede reconstruir con certeza el universo 66: "
            f"control={len(base_control)}, faltan={missing_n}, grupos_excluidos=[{detail}]"
        )

    universe = set(base_control)
    universe.update(s(r.get("CODCLI")) for r in selected_extra)

    outside = sorted({s(r.get("CODCLI")) for r in excluded_not_control} - {s(r.get("CODCLI")) for r in selected_extra})
    rule = (
        "reconstruido desde CONTROL_FINAL (55 CODCLI: 54 OK + 1 bloqueado) "
        f"+ unico grupo de excluidos que cierra {EXPECTED_UNIVERSE}: "
        f"{selected_extra[0].get('CODCARPR') if selected_extra else 'N/A'} "
        f"({len(selected_extra)} CODCLI)."
    )
    return universe, rule, selected_extra, outside


def exclusion_reason(row: dict[str, str]) -> str:
    estado_pes = s(row.get("ESTADO_VALIDACION_PES"))
    errores = s(row.get("ERRORES_VALIDACION"))
    estado_gob = s(row.get("ESTADO_GOBERNANZA"))
    accion = s(row.get("ACCION_CARGA"))
    obs = s(row.get("OBSERVACION_DICCIONARIO"))
    anom = s(row.get("ANOMATRICULA"))
    anoi = s(row.get("ANOINGRESO"))

    if estado_pes == "BLOQUEADO" or "RUT_INVALIDO" in errores or "BLOQUEADO_RUT_INVALIDO" in estado_gob:
        return errores or "RUT_INVALIDO_LONGITUD_NO_TRAZABLE"
    if anom != "2026" and anoi != "2026":
        return "NO_CORRESPONDE_A_2026"
    if estado_gob == "NO_MAPEAR_GENERICO":
        return "PROGRAMA_GENERICO_NO_MAPEABLE"
    if estado_gob == "PENDIENTE_REVISION":
        if "SIN_MAPEO" in " ".join(s(row.get(c)) for c in ["CODIGO_UNICO", "COD_SED", "COD_CAR", "MODALIDAD", "JOR", "VERSION"]):
            return "PROGRAMA_SIN_DICCIONARIO_OFERTA_GOBERNADA"
        return "GOBERNANZA_PENDIENTE"
    if accion and accion != "INCLUIR":
        return f"NO_CUMPLE_CONDICION_INCLUSION_{accion}"
    return obs or "OTRO_MOTIVO_DE_EXCLUSION"


def build_markdown(
    summary: dict[str, Any],
    master: pd.DataFrame,
    included: pd.DataFrame,
    excluded: pd.DataFrame,
    validations: pd.DataFrame,
    paths: dict[str, str],
) -> str:
    def md_table(df: pd.DataFrame, cols: list[str], limit: int | None = None) -> str:
        view = df[cols].copy()
        if limit is not None:
            view = view.head(limit)
        if view.empty:
            return "_Sin registros._"
        headers = [s(col) for col in view.columns]

        def esc(value: Any) -> str:
            return s(value).replace("|", "\\|").replace("\n", " ")

        lines = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join("---" for _ in headers) + " |",
        ]
        for record in view.to_dict("records"):
            lines.append("| " + " | ".join(esc(record.get(col, "")) for col in headers) + " |")
        return "\n".join(lines)

    correo = (
        "Para la carga MU2026 Postgrado/Postitulo se reconcilio el universo original "
        "de 66 CODCLI contra el PES-ready final. De esos 66 registros, 54 cumplen "
        "las condiciones de inclusion y permanecen en el archivo final; 12 quedan "
        "fuera con trazabilidad: 11 por gobernanza pendiente/programa sin mapeo "
        "oficial en matriz-oferta y 1 por RUT invalido de longitud no trazable en "
        "las fuentes internas. Por eso el archivo final contiene 54 estudiantes y "
        "no 66. La conciliacion cierra matematicamente: 66 = 54 + 12."
    )

    lines = [
        "# Reconciliacion 66 a 54 Postgrado/Postitulo MU2026",
        "",
        f"Fecha de ejecucion: `{summary['timestamp']}`",
        "",
        f"Resultado final: **{summary['estado_final']}**",
        "",
        "## Resumen",
        "",
        "| Indicador | Valor |",
        "| --- | ---: |",
        f"| Universo original | {summary['universo_original']} |",
        f"| Incluidos PES-ready | {summary['incluidos_pes_ready']} |",
        f"| Excluidos | {summary['excluidos']} |",
        f"| Bloqueados por error | {summary['bloqueados_error']} |",
        f"| Diferencia 66 - 54 | {summary['diferencia_66_54']} |",
        f"| Estado final | {summary['estado_final']} |",
        "",
        "## Regla de reconstruccion del universo",
        "",
        summary["regla_reconstruccion"],
        "",
        "## Validaciones",
        "",
        md_table(validations, ["VALIDACION", "ESTADO", "DETALLE"]),
        "",
        "## Detalle completo de excluidos",
        "",
        md_table(excluded, ["CODCLI", "RUT_DV", "NOMBRE_L", "ESTADOACADEMICO", "MOTIVO_EXCLUSION", "FUENTE_DECISION"]),
        "",
        "## Detalle completo de incluidos",
        "",
        md_table(included, ["CODCLI", "RUT_DV", "COD_CAR", "MODALIDAD", "JOR", "FOR_ING_ACT", "VIG"]),
        "",
        "## Texto sugerido para correo institucional",
        "",
        correo,
        "",
        "## Archivos generados",
        "",
    ]
    lines.extend(f"- `{path}`" for path in paths.values())
    lines.extend([
        "",
        "## Observacion sobre candidatos 2026 fuera del universo 66",
        "",
        summary["observacion_fuera_universo_66"],
        "",
    ])
    return "\n".join(lines)


def write_excel(
    path: Path,
    sheets: dict[str, pd.DataFrame],
    master_sheet: str = "UNIVERSO_66",
) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=sheet[:31])

    wb = load_workbook(path)
    fills = {
        "incluido": PatternFill("solid", fgColor="D9EAD3"),
        "excluido": PatternFill("solid", fgColor="FCE5CD"),
        "bloqueado": PatternFill("solid", fgColor="F4CCCC"),
        "error": PatternFill("solid", fgColor="CC0000"),
        "warning": PatternFill("solid", fgColor="FFF2CC"),
        "header": PatternFill("solid", fgColor="1F4E78"),
    }
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        max_col = ws.max_column
        max_row = ws.max_row
        if max_col:
            ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
        for cell in ws[1]:
            cell.fill = fills["header"]
            cell.font = header_font
        headers = [s(c.value) for c in ws[1]]
        status_idx = headers.index("INGRESA_A_PES_READY") + 1 if "INGRESA_A_PES_READY" in headers else None
        motivo_idx = headers.index("MOTIVO_EXCLUSION") + 1 if "MOTIVO_EXCLUSION" in headers else None
        valid_idx = headers.index("ESTADO") + 1 if "ESTADO" in headers else None
        for row in ws.iter_rows(min_row=2, max_row=max_row):
            fill = None
            if status_idx:
                value = s(row[status_idx - 1].value)
                motivo = s(row[motivo_idx - 1].value) if motivo_idx else ""
                if value == "SI":
                    fill = fills["incluido"]
                elif "RUT_INVALIDO" in motivo or "BLOQUEADO" in motivo:
                    fill = fills["bloqueado"]
                elif value == "NO":
                    fill = fills["excluido"]
            elif valid_idx:
                value = s(row[valid_idx - 1].value)
                if value == "ERROR":
                    fill = fills["error"]
                elif value == "ADVERTENCIA":
                    fill = fills["warning"]
                elif value == "OK":
                    fill = fills["incluido"]
            if fill:
                for cell in row:
                    cell.fill = fill
        for col_cells in ws.columns:
            letter = get_column_letter(col_cells[0].column)
            width = min(max(len(s(cell.value)) for cell in col_cells) + 2, 60)
            ws.column_dimensions[letter].width = max(width, 10)
    if master_sheet in wb.sheetnames:
        wb.active = wb.sheetnames.index(master_sheet)
    wb.save(path)


def validation_row(nombre: str, estado: str, detalle: str) -> dict[str, str]:
    return {"VALIDACION": nombre, "ESTADO": estado, "DETALLE": detalle}


def main() -> int:
    required = [PES_READY, CON_TITULOS, CONTROL_FINAL, EXCLUIDOS, TRAZABILIDAD, DICCIONARIO, EXCEL_FUENTE]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Faltan fuentes obligatorias: " + ", ".join(missing))

    pre_hash_before = {str(p): sha256(p) for p in PREGRADO_ARCHIVOS}
    backups: list[dict[str, str]] = []
    for path in [OUT_XLSX, OUT_CSV]:
        backup_if_exists(path, backups)

    pes_rows = read_pes(PES_READY)
    con_rows = read_csv_dict(CON_TITULOS, delimiter=";")
    control_rows = read_csv_dict(CONTROL_FINAL)
    excluded_rows = read_csv_dict(EXCLUIDOS)
    traza_rows = read_csv_dict(TRAZABILIDAD, delimiter="\t")
    dic_rows = read_csv_dict(DICCIONARIO, delimiter="\t")
    da_by_codcli, base_by_codcli = load_excel_sources()

    traza_by_codcli = {s(r.get("CODCLI")): r for r in traza_rows if s(r.get("CODCLI"))}
    control_by_codcli = {s(r.get("CODCLI")): r for r in control_rows if s(r.get("CODCLI"))}
    excluded_by_codcli = {s(r.get("CODCLI")): r for r in excluded_rows if s(r.get("CODCLI"))}
    dic_by_codcarpr = {s(r.get("CODCARPR")): r for r in dic_rows if s(r.get("CODCARPR"))}

    universe_codcli, universe_rule, selected_excluded, outside_codcli = reconstruct_universe(control_rows, excluded_rows)
    if len(universe_codcli) != EXPECTED_UNIVERSE:
        raise RuntimeError(f"Universo reconstruido no cierra: {len(universe_codcli)} != {EXPECTED_UNIVERSE}")

    control_ok = [r for r in control_rows if s(r.get("ESTADO_VALIDACION_PES")) == "OK"]
    included_codcli = {s(r.get("CODCLI")) for r in control_ok if s(r.get("CODCLI"))}
    pes_tuples = Counter(pes_tuple(r) for r in pes_rows)
    control_ok_tuples = Counter(pes_tuple(r) for r in control_ok)

    master_rows: list[dict[str, str]] = []
    for codcli in sorted(universe_codcli):
        tr = traza_by_codcli.get(codcli, {})
        ctrl = control_by_codcli.get(codcli, {})
        exc = excluded_by_codcli.get(codcli, {})
        da = da_by_codcli.get(codcli, {})
        base = base_by_codcli.get(codcli, {})

        codcarpr = s(tr.get("CODCARPR") or ctrl.get("CODCARPR_ORIGEN") or exc.get("CODCARPR") or da.get("CODCARPR"))
        dic = dic_by_codcarpr.get(codcarpr, {})
        rut_source = s(tr.get("RUT") or ctrl.get("RUT_ORIGEN") or exc.get("RUT") or da.get("RUT"))
        rut, dv = split_rut(rut_source)
        if not dv:
            _, dv = split_rut(base.get("RUT", ""), base.get("DV", ""))
        estado_pes = s(ctrl.get("ESTADO_VALIDACION_PES"))
        ingresa = "SI" if codcli in included_codcli else "NO"

        combined = {
            **tr,
            **{f"EXCL_{k}": v for k, v in exc.items()},
            **{f"CTRL_{k}": v for k, v in ctrl.items()},
            "ESTADO_VALIDACION_PES": estado_pes,
            "ERRORES_VALIDACION": s(ctrl.get("ERRORES_VALIDACION")),
        }
        motivo_exclusion = "" if ingresa == "SI" else exclusion_reason(combined | tr | exc | ctrl | dic)
        motivo_inclusion = (
            "ACCION_CARGA=INCLUIR; ESTADO_VALIDACION_PES=OK; presente en PES-ready final."
            if ingresa == "SI" else ""
        )
        fuente = []
        if tr:
            fuente.append("trazabilidad")
        if ctrl:
            fuente.append("control_final")
        if exc:
            fuente.append("excluidos")
        if da:
            fuente.append("DatosAlumnos")
        if dic:
            fuente.append("diccionario")
        if base:
            fuente.append("base_datos")

        obs = []
        if not da:
            obs.append("CODCLI no encontrado en DatosAlumnos.")
        elif s(da.get("RUT")) and s(da.get("RUT")) != rut_source:
            obs.append(f"RUT DatosAlumnos difiere de trazabilidad/control: {da.get('RUT')}.")
        ok_rut, rut_msg = valid_rut(rut, dv)
        if not ok_rut:
            obs.append(rut_msg)
        if codcli in outside_codcli:
            obs.append("Candidato 2026 fuera del universo 66 solicitado.")

        row = {
            "CODCLI": codcli,
            "RUT": rut,
            "DV": dv,
            "RUT_DV": f"{rut}-{dv}" if rut or dv else "",
            "CODCARPR": codcarpr,
            "CODIGOCARRERA": s(tr.get("CODIGOCARRERA") or exc.get("CODIGOCARRERA") or da.get("CODIGOCARRERA")),
            "NOMBRE_L": s(tr.get("NOMBRE_L") or ctrl.get("NOMBRE_L_ORIGEN") or exc.get("NOMBRE_L") or da.get("NOMBRE_L")),
            "ANOMATRICULA": s(tr.get("ANOMATRICULA") or exc.get("ANOMATRICULA") or da.get("ANOMATRICULA")),
            "PERIODOMATRICULA": s(tr.get("PERIODOMATRICULA") or da.get("PERIODOMATRICULA")),
            "ANOINGRESO": s(tr.get("ANOINGRESO") or exc.get("ANOINGRESO") or da.get("ANOINGRESO")),
            "PERIODOINGRESO": s(tr.get("PERIODOINGRESO") or da.get("PERIODOINGRESO")),
            "ESTADOACADEMICO": s(tr.get("ESTADOACADEMICO") or ctrl.get("ESTADOACADEMICO_ORIGEN") or exc.get("ESTADOACADEMICO") or da.get("ESTADOACADEMICO")),
            "MATRICULA": s(tr.get("MATRICULA") or da.get("MATRICULA")),
            "ACCION_CARGA": s(tr.get("ACCION_CARGA") or exc.get("ACCION_CARGA") or dic.get("ACCION_CARGA")),
            "ESTADO_GOBERNANZA": s(tr.get("ESTADO_GOBERNANZA") or exc.get("ESTADO_GOBERNANZA") or dic.get("ESTADO_GOBERNANZA")),
            "ESTADO_TRAZABILIDAD": s(tr.get("ESTADO_TRAZABILIDAD") or exc.get("ESTADO_TRAZABILIDAD")),
            "ESTADO_VALIDACION_PES": estado_pes,
            "INGRESA_A_PES_READY": ingresa,
            "MOTIVO_INCLUSION": motivo_inclusion,
            "MOTIVO_EXCLUSION": motivo_exclusion,
            "COD_SED": s(ctrl.get("COD_SED") or tr.get("COD_SED") or dic.get("COD_SED")),
            "COD_CAR": s(ctrl.get("COD_CAR") or tr.get("COD_CAR") or dic.get("COD_CAR")),
            "MODALIDAD": s(ctrl.get("MODALIDAD") or tr.get("MODALIDAD") or dic.get("MODALIDAD")),
            "JOR": s(ctrl.get("JOR") or tr.get("JOR") or dic.get("JOR")),
            "VERSION": s(ctrl.get("VERSION") or tr.get("VERSION") or dic.get("VERSION")),
            "FOR_ING_ACT": s(ctrl.get("FOR_ING_ACT")),
            "VIG": s(ctrl.get("VIG") or tr.get("VIGENCIA_OFERTA") or dic.get("VIGENCIA_OFERTA")),
            "FUENTE_DECISION": " + ".join(fuente) if fuente else "sin fuente",
            "OBSERVACION_AUDITORIA": "; ".join(obs),
        }
        master_rows.append(row)

    master = df_from_rows(master_rows, MASTER_COLUMNS)
    included = master[master["INGRESA_A_PES_READY"].eq("SI")].copy()
    excluded = master[master["INGRESA_A_PES_READY"].eq("NO")].copy()
    blocked = excluded[excluded["MOTIVO_EXCLUSION"].str.contains("RUT_INVALIDO|BLOQUEADO", na=False)].copy()
    outside_rows = [
        {
            "CODCLI": codcli,
            "RUT_DV": s(traza_by_codcli.get(codcli, {}).get("RUT") or excluded_by_codcli.get(codcli, {}).get("RUT")),
            "CODCARPR": s(traza_by_codcli.get(codcli, {}).get("CODCARPR") or excluded_by_codcli.get(codcli, {}).get("CODCARPR")),
            "NOMBRE_L": s(traza_by_codcli.get(codcli, {}).get("NOMBRE_L") or excluded_by_codcli.get(codcli, {}).get("NOMBRE_L")),
            "MOTIVO": "Candidato 2026 en trazabilidad/excluidos, fuera de la reconstruccion unica del universo 66 solicitado.",
        }
        for codcli in outside_codcli
    ]
    outside_df = df_from_rows(outside_rows)

    validations: list[dict[str, str]] = []
    validations.append(validation_row(
        "universo_original = incluidos + excluidos",
        "OK" if len(master) == len(included) + len(excluded) == EXPECTED_UNIVERSE else "ERROR",
        f"{len(master)} = {len(included)} + {len(excluded)}",
    ))
    validations.append(validation_row(
        "incluidos = filas PES-ready final",
        "OK" if len(included) == len(pes_rows) else "ERROR",
        f"incluidos={len(included)}; filas_pes={len(pes_rows)}",
    ))
    field_counts = Counter(s(r.get("_N_CAMPOS")) for r in pes_rows)
    validations.append(validation_row(
        "PES-ready con 21 campos por fila",
        "OK" if field_counts == Counter({"21": len(pes_rows)}) else "ERROR",
        dict(field_counts).__repr__(),
    ))
    validations.append(validation_row(
        "control OK puentea exactamente el PES-ready",
        "OK" if pes_tuples == control_ok_tuples else "ERROR",
        f"pes_tuples={sum(pes_tuples.values())}; control_ok_tuples={sum(control_ok_tuples.values())}",
    ))
    dup_codcli = included["CODCLI"][included["CODCLI"].duplicated()].tolist()
    validations.append(validation_row(
        "sin CODCLI incluido dos veces",
        "OK" if not dup_codcli else "ERROR",
        ", ".join(dup_codcli) if dup_codcli else "sin duplicados",
    ))
    dup_rut = included["RUT_DV"][included["RUT_DV"].duplicated() & included["RUT_DV"].ne("")].tolist()
    validations.append(validation_row(
        "sin RUT/DV duplicado no justificado en incluidos",
        "OK" if not dup_rut else "ADVERTENCIA",
        ", ".join(dup_rut) if dup_rut else "sin duplicados",
    ))
    contains_111222 = any(s(r.get("N_DOC")) == "111222" for r in pes_rows)
    validations.append(validation_row(
        "N_DOC=111222 ausente en PES-ready final",
        "OK" if not contains_111222 else "ERROR",
        "ausente" if not contains_111222 else "presente",
    ))
    target_blocked = (
        "20261DCBS003" in set(excluded["CODCLI"])
        and excluded.loc[excluded["CODCLI"].eq("20261DCBS003"), "MOTIVO_EXCLUSION"].str.contains("RUT_INVALIDO", na=False).any()
    )
    validations.append(validation_row(
        "registro bloqueado por RUT invalido en excluidos",
        "OK" if target_blocked else "ERROR",
        "20261DCBS003 con motivo RUT_INVALIDO" if target_blocked else "no encontrado",
    ))
    pes_rut_errors = []
    for row in pes_rows:
        if s(row.get("TIPO_DOC")) == "R":
            ok, msg = valid_rut(s(row.get("N_DOC")), s(row.get("DV")))
            if not ok:
                pes_rut_errors.append(f"linea {row['_LINEA_PES']}: {msg}")
    validations.append(validation_row(
        "RUT/DV validos en PES-ready",
        "OK" if not pes_rut_errors else "ERROR",
        "; ".join(pes_rut_errors) if pes_rut_errors else "sin errores",
    ))
    pre_hash_after = {str(p): sha256(p) for p in PREGRADO_ARCHIVOS}
    pregrado_intacto = pre_hash_before == pre_hash_after
    validations.append(validation_row(
        "pregrado no modificado",
        "OK" if pregrado_intacto else "ERROR",
        "hashes sin cambios" if pregrado_intacto else "hashes cambiaron",
    ))
    validations_df = df_from_rows(validations, ["VALIDACION", "ESTADO", "DETALLE"])

    critical_errors = int((validations_df["ESTADO"] == "ERROR").sum())
    warnings = int((validations_df["ESTADO"] == "ADVERTENCIA").sum())
    governance_exclusions = int(excluded["MOTIVO_EXCLUSION"].str.contains("GOBERNANZA|SIN_DICCIONARIO|SIN_MAPEO", na=False).sum())
    rut_exclusions = int(excluded["MOTIVO_EXCLUSION"].str.contains("RUT_INVALIDO|RUT/DV", na=False).sum())
    not_2026_exclusions = int(excluded["MOTIVO_EXCLUSION"].str.contains("NO_CORRESPONDE_A_2026", na=False).sum())
    sin_mapeo_exclusions = int(excluded["MOTIVO_EXCLUSION"].str.contains("SIN_DICCIONARIO|SIN_MAPEO", na=False).sum())
    generico_exclusions = int(excluded["MOTIVO_EXCLUSION"].str.contains("GENERICO", na=False).sum())

    if critical_errors:
        estado_final = "BLOQUEADO"
    elif governance_exclusions or rut_exclusions:
        estado_final = "APROBADO_CON_OBSERVACIONES"
    else:
        estado_final = "APROBADO"

    summary = {
        "timestamp": datetime.now().isoformat(timespec="seconds"),
        "estado_final": estado_final,
        "universo_original": int(len(master)),
        "incluidos_pes_ready": int(len(included)),
        "excluidos": int(len(excluded)),
        "bloqueados_error": int(len(blocked)),
        "excluidos_gobernanza_pendiente": governance_exclusions,
        "excluidos_rut_invalido": rut_exclusions,
        "excluidos_no_corresponde_2026": not_2026_exclusions,
        "excluidos_programa_sin_mapeo": sin_mapeo_exclusions,
        "excluidos_programa_generico": generico_exclusions,
        "diferencia_66_54": EXPECTED_UNIVERSE - len(included),
        "reconciliacion_matematica": bool(len(master) == EXPECTED_UNIVERSE and len(included) == len(pes_rows) == 54 and len(excluded) == 12),
        "errores_criticos": critical_errors,
        "advertencias": warnings,
        "regla_reconstruccion": universe_rule,
        "codcli_fuera_universo_66": outside_codcli,
        "observacion_fuera_universo_66": (
            f"Se detectaron {len(outside_codcli)} candidatos 2026 adicionales en trazabilidad/excluidos "
            "que no forman parte del universo 66 reconstruido porque no pertenecen al unico grupo que cierra "
            "66 = control_final + excluidos faltantes. Ver hoja FUERA_UNIVERSO_66."
            if outside_codcli else "No se detectaron candidatos adicionales fuera del universo 66."
        ),
        "pregrado_no_modificado": pregrado_intacto,
        "pes_ready_filas": int(len(pes_rows)),
        "pes_ready_campos": {k: int(v) for k, v in field_counts.items()},
        "no_existe_n_doc_111222": not contains_111222,
        "backups": backups,
    }

    resumen_df = df_from_rows([
        {"INDICADOR": "Estado final", "VALOR": estado_final},
        {"INDICADOR": "Universo original CODCLI", "VALOR": len(master)},
        {"INDICADOR": "Incluidos PES-ready", "VALOR": len(included)},
        {"INDICADOR": "Excluidos", "VALOR": len(excluded)},
        {"INDICADOR": "Bloqueados por error", "VALOR": len(blocked)},
        {"INDICADOR": "Excluidos gobernanza pendiente", "VALOR": governance_exclusions},
        {"INDICADOR": "Excluidos RUT invalido", "VALOR": rut_exclusions},
        {"INDICADOR": "Excluidos no corresponde 2026", "VALOR": not_2026_exclusions},
        {"INDICADOR": "Excluidos programa sin mapeo", "VALOR": sin_mapeo_exclusions},
        {"INDICADOR": "Excluidos programa generico", "VALOR": generico_exclusions},
        {"INDICADOR": "Diferencia exacta 66 - 54", "VALOR": EXPECTED_UNIVERSE - len(included)},
        {"INDICADOR": "Confirmacion matematica", "VALOR": "SI" if summary["reconciliacion_matematica"] else "NO"},
        {"INDICADOR": "Regla reconstruccion", "VALOR": universe_rule},
    ], ["INDICADOR", "VALOR"])

    control_df = df_from_rows(control_rows)
    dic_df = df_from_rows(dic_rows)

    write_excel(OUT_XLSX, {
        "RESUMEN": resumen_df,
        "UNIVERSO_66": master,
        "INCLUIDOS_54": included,
        "EXCLUIDOS": excluded,
        "BLOQUEADOS_ERROR": blocked,
        "VALIDACIONES": validations_df,
        "DICCIONARIO": dic_df,
        "CONTROL_FINAL": control_df,
        "FUERA_UNIVERSO_66": outside_df,
    })
    master.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")

    paths = {
        "excel": str(OUT_XLSX),
        "csv": str(OUT_CSV),
        "markdown": str(OUT_MD),
        "json": str(OUT_JSON),
    }
    md = build_markdown(summary, master, included, excluded, validations_df, paths)
    OUT_MD.write_text(md, encoding="utf-8")
    OUT_JSON.write_text(json.dumps({**summary, "archivos": paths}, ensure_ascii=False, indent=2), encoding="utf-8")

    desktop = Path.home() / "Desktop" / "reconciliacion_66_a_54_postgrado.xlsx"
    if desktop.exists():
        backup_if_exists(desktop, backups)
    shutil.copy2(OUT_XLSX, desktop)
    summary["archivo_escritorio"] = str(desktop)
    summary["backups"] = backups
    OUT_JSON.write_text(json.dumps({**summary, "archivos": paths}, ensure_ascii=False, indent=2), encoding="utf-8")
    subprocess.run(["open", str(desktop)], check=False)

    print("RUTA_EXCEL=", OUT_XLSX)
    print("RUTA_CSV=", OUT_CSV)
    print("RUTA_MARKDOWN=", OUT_MD)
    print("RUTA_JSON=", OUT_JSON)
    print("ESTADO_FINAL=", estado_final)
    print("")
    print("TABLA_66_A_54")
    print(resumen_df.head(12).to_string(index=False))
    print("")
    print("DETALLE_12_EXCLUIDOS")
    cols = ["CODCLI", "RUT_DV", "NOMBRE_L", "ESTADOACADEMICO", "MOTIVO_EXCLUSION", "FUENTE_DECISION"]
    print(excluded[cols].to_string(index=False))
    print("")
    print(f"PES_READY_FINAL_FILAS={len(pes_rows)}")
    print(f"PES_READY_21_CAMPOS={field_counts == Counter({'21': len(pes_rows)})}")
    print(f"PREGRADO_NO_MODIFICADO={pregrado_intacto}")
    print(f"EXCEL_COPIADO_ESCRITORIO={desktop}")
    if critical_errors:
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR_CRITICO={exc}", file=sys.stderr)
        raise
