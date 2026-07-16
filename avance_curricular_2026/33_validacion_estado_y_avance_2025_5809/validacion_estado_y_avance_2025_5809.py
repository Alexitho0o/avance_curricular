#!/usr/bin/env python3
# -*- coding: utf-8 -*-
from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import sys
import unicodedata
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore", message="Title is more than 31 characters")

PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Validacion estado academico + avance academico 2025 - 5809"
ANIO_REFERENCIA = 2025
REPO_ROOT = Path("/Users/alexi/Documents/GitHub/avance_curricular").resolve()
OUTPUT_DIR = REPO_ROOT / "avance_curricular_2026/33_validacion_estado_y_avance_2025_5809"
PROMEDIOS = REPO_ROOT / "avance_curricular_2026/25_actualizacion_fuente_promedios/CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/00_FUENTE_CONGELADA/PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
CSV_5809 = REPO_ROOT / "avance_curricular_2026/14_sies_ready_5809/SIES_READY_TECNICO_5809_22_COLUMNAS_20260704_151418/01_ARCHIVO_SIES_READY_TECNICO/5809_MATRICULA_AVANCE_CURRICULAR_2026_SIES_READY_TECNICO.csv"
AUDITORIA_CORREGIDA = REPO_ROOT / "avance_curricular_2026/29_auditoria_corregida_logica_2025_5809/AUDITORIA_CORREGIDA_LOGICA_2025_5809.xlsx"
MAPEO_RESOLUCION = REPO_ROOT / "avance_curricular_2026/30_mapeo_resolucion_bloqueos_5809_2025/MAPEO_RESOLUCION_BLOQUEOS_5809_2025_20260708_092026/02_RESULTADOS/MAPEO_RESOLUCION_BLOQUEOS_5809_LOGICA_2025.xlsx"
EXCEL_OUT = OUTPUT_DIR / "VALIDACION_ESTADO_Y_AVANCE_2025_5809.xlsx"
MD_OUT = OUTPUT_DIR / "INFORME_VALIDACION_ESTADO_Y_AVANCE_2025_5809.md"
MANIFEST_OUT = OUTPUT_DIR / "manifest_validacion_estado_y_avance_2025_5809.json"

CSV_COLUMNS = [
    "CODIGO_IES_NUM",
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
    "VIGENCIA",
]


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.upper()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def norm_col(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", norm_text(value))


def clean_key(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    if re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def to_int(value: Any) -> int | None:
    text = clean_key(value)
    if text == "":
        return None
    try:
        return int(float(text.replace(",", ".")))
    except Exception:
        return None


def numeric_series(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series.astype(str).str.replace(",", ".", regex=False).str.strip(), errors="coerce")


def normalize_rut(num: Any, dv: Any = None) -> str:
    raw = clean_key(num).upper().replace(".", "").replace(" ", "")
    raw_dv = clean_key(dv).upper().replace(".", "").replace(" ", "")
    if not raw:
        return ""
    if "-" in raw:
        left, right = raw.rsplit("-", 1)
        raw = left
        raw_dv = right
    body = re.sub(r"[^0-9]", "", raw)
    digit = re.sub(r"[^0-9K]", "", raw_dv)
    if not body:
        return ""
    if digit:
        return f"{body}-{digit[-1]}"
    return body


def normalize_si_no(value: Any) -> str:
    text = norm_text(value)
    if text in {"SI", "S", "1", "TRUE", "VERDADERO"}:
        return "SI"
    if text in {"NO", "N", "0", "FALSE", "FALSO"}:
        return "NO"
    return text


def sha256_file(path: Path) -> str:
    if not path.exists():
        return "NO_EXISTE"
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def json_safe(value: Any) -> Any:
    if isinstance(value, Path):
        return str(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return str(value)
    if isinstance(value, set):
        return sorted(value)
    return value


def safe_join(values: list[Any] | set[Any], sep: str = " | ") -> str:
    cleaned = []
    for value in values:
        text = clean_key(value)
        if text and text not in cleaned:
            cleaned.append(text)
    return sep.join(cleaned)


def classify_estado_asignatura(value: Any) -> tuple[str, str]:
    text = norm_text(value)
    if not text:
        return "NO_CLASIFICADO", "Valor vacio"
    no_hits = []
    yes_hits = []
    if "REPROB" in text:
        no_hits.append("contiene REPROB")
    if re.search(r"\bREP\b", text):
        no_hits.append("contiene REP")
    if re.search(r"\bNO\s+APROB", text):
        no_hits.append("contiene NO APROB")
    if re.search(r"\bNCR\b", text):
        no_hits.append("contiene NCR")
    if re.search(r"(?<!NO )APROB", text):
        yes_hits.append("contiene APROB")
    if re.search(r"\bAPR\b", text):
        yes_hits.append("contiene APR")
    if "CONVALID" in text:
        yes_hits.append("contiene CONVALID")
    if "HOMOLOG" in text:
        yes_hits.append("contiene HOMOLOG")
    if "RECONOC" in text:
        yes_hits.append("contiene RECONOC")
    if yes_hits and no_hits:
        return "NO_CLASIFICADO", "Senales contradictorias: " + "; ".join(yes_hits + no_hits)
    if yes_hits:
        return "APROBADO", "; ".join(yes_hits)
    if no_hits:
        return "NO_APROBADO", "; ".join(no_hits)
    return "NO_CLASIFICADO", "Sin patron explicito permitido"


def classify_estado_academico(value: Any) -> tuple[str, str]:
    text = norm_text(value)
    if not text:
        return "SIN_ESTADO", "Valor vacio"
    if re.search(r"\bELIMINAD[OA]\b", text):
        return "EXPLICATIVO_AUSENCIA_2025", "Contiene ELIMINADO/ELIMINADA"
    if "SIN ACTIVIDAD ACADEMICA" in text:
        return "EXPLICATIVO_AUSENCIA_2025", "Contiene SIN ACTIVIDAD ACADEMICA"
    return "OBSERVADO_NO_EXPLICATIVO", "Estado observado; no explica ausencia anual por si solo"


def is_estado_explicativo(value: Any) -> bool:
    return classify_estado_academico(value)[0] == "EXPLICATIVO_AUSENCIA_2025"


def detect_col(df: pd.DataFrame, preferred: list[str], contains: list[str] | None = None, exclude: list[str] | None = None) -> str | None:
    if df.empty:
        return None
    contains = contains or []
    exclude = exclude or []
    col_map = {norm_col(c): c for c in df.columns}
    for candidate in preferred:
        key = norm_col(candidate)
        if key in col_map:
            return str(col_map[key])
    for col in df.columns:
        ncol = norm_col(col)
        if any(token in ncol for token in exclude):
            continue
        if any(norm_col(candidate) in ncol for candidate in preferred):
            return str(col)
    for col in df.columns:
        ncol = norm_col(col)
        if any(token in ncol for token in exclude):
            continue
        if any(token in ncol for token in contains):
            return str(col)
    return None


def candidate_cols(df: pd.DataFrame, tokens: list[str], exclude: list[str] | None = None) -> list[str]:
    exclude = exclude or []
    out = []
    for col in df.columns:
        ncol = norm_col(col)
        if any(ex in ncol for ex in exclude):
            continue
        if any(token in ncol for token in tokens):
            out.append(str(col))
    return out


def detect_mapping_cols(df: pd.DataFrame) -> dict[str, str | None]:
    return {
        "CODCLI": detect_col(df, ["CODCLI", "COD_CLIENTE", "CODALUMNO"], contains=["CODCLI"]),
        "RUT": detect_col(df, ["RUT", "RUN", "NUM_DOCUMENTO", "NUMDOCUMENTO", "DOCUMENTO"], contains=["RUT"], exclude=["RESPONSABLE", "EJECUTIVO"]),
        "DV": detect_col(df, ["DV", "DIG", "DIGITO", "DIGITO_VERIFICADOR", "DVRUT"], contains=["DIGITOVERIFICADOR"]),
    }


def detect_hoja1_columns(hoja1: pd.DataFrame) -> tuple[dict[str, Any], pd.DataFrame]:
    selected: dict[str, Any] = {}
    selected["CODCLI"] = detect_col(hoja1, ["CODCLI"], contains=["CODCLI"])
    selected["RUT"] = detect_col(hoja1, ["RUT"], contains=["RUT"], exclude=["RESPONSABLE", "EJECUTIVO"])
    selected["DV"] = detect_col(hoja1, ["DIG", "DV"], contains=["DIGITOVERIFICADOR"])
    selected["ANO"] = detect_col(hoja1, ["ANO", "ANIO"], contains=["ANO"], exclude=["INGRESO", "MATRICULA", "EGRESO"])
    selected["PERIODO"] = detect_col(hoja1, ["PERIODO"], contains=["PERIODO", "SEMESTRE"], exclude=["INGRESO", "MATRICULA"])
    course_cols = []
    for name in ["CODRAMO", "RAMOEQUIV", "ASIGNATURA"]:
        col = detect_col(hoja1, [name], contains=[norm_col(name)])
        if col and col not in course_cols:
            course_cols.append(col)
    selected["CODRAMO_ASIGNATURA"] = course_cols[0] if course_cols else None
    selected["COURSE_FALLBACK_COLS"] = course_cols
    estado_candidates = candidate_cols(hoja1, ["ESTADO"], exclude=["ESTADOACADEMICO", "ACADEMICO", "MATRICULA"])
    if "DESCRIPCION_ESTADO" in hoja1.columns and "DESCRIPCION_ESTADO" not in estado_candidates:
        estado_candidates.append("DESCRIPCION_ESTADO")
    rows = []
    best_col = None
    best_score = -1.0
    for col in estado_candidates:
        values = hoja1[col].fillna("").map(lambda x: classify_estado_asignatura(x)[0])
        nonempty = hoja1[col].fillna("").map(lambda x: norm_text(x) != "")
        denom = int(nonempty.sum())
        score = float((values[nonempty] != "NO_CLASIFICADO").mean()) if denom else 0.0
        rows.append({
            "tipo_registro": "EVALUACION_ESTADO_ASIGNATURA",
            "hoja": "Hoja1",
            "columna": col,
            "pct_clasificable_patrones_explicitos": round(score, 6),
            "seleccionada": "NO",
            "motivo": "Evaluacion de patrones APROB/APR/CONVALID/HOMOLOG/RECONOC y REPROB/REP/NO APROB/NCR",
        })
        if score > best_score:
            best_score = score
            best_col = col
    selected["ESTADO_ASIGNATURA"] = best_col
    for row in rows:
        if row["columna"] == best_col:
            row["seleccionada"] = "SI"
    return selected, pd.DataFrame(rows)


def build_column_diagnostics(sheets: dict[str, pd.DataFrame], selected: dict[str, Any], estado_eval: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for sheet, df in sheets.items():
        cols = [str(c) for c in df.columns]
        rows.append({
            "tipo_registro": "HOJA_PROMEDIOS",
            "hoja": sheet,
            "cantidad_filas": len(df),
            "cantidad_columnas": len(df.columns),
            "columnas": safe_join(cols),
            "candidatas_CODCLI": safe_join(candidate_cols(df, ["CODCLI"])),
            "candidatas_RUT": safe_join(candidate_cols(df, ["RUT"], exclude=["RESPONSABLE", "EJECUTIVO"])),
            "candidatas_DV": safe_join(candidate_cols(df, ["DV", "DIG"])),
            "candidatas_ANO": safe_join(candidate_cols(df, ["ANO", "ANIO"], exclude=["INGRESO", "MATRICULA"])),
            "candidatas_PERIODO": safe_join(candidate_cols(df, ["PERIODO", "SEMESTRE"])),
            "candidatas_CODRAMO_ASIGNATURA": safe_join(candidate_cols(df, ["CODRAMO", "RAMOEQUIV", "ASIGNATURA"])),
            "candidatas_ESTADO_ASIGNATURA": safe_join(candidate_cols(df, ["ESTADO"], exclude=["ESTADOACADEMICO", "ACADEMICO"])),
        })
    for logical, col in selected.items():
        if logical == "COURSE_FALLBACK_COLS":
            continue
        rows.append({
            "tipo_registro": "SELECCION_TECNICA_HOJA1",
            "hoja": "Hoja1",
            "campo_logico": logical,
            "columna_seleccionada": col or "",
            "motivo": "Seleccion para recalculo anual 2025; no mezcla estado academico con estado de asignatura",
        })
    if not estado_eval.empty:
        return pd.concat([pd.DataFrame(rows), estado_eval], ignore_index=True)
    return pd.DataFrame(rows)


def prepare_hoja1(hoja1: pd.DataFrame, selected: dict[str, Any]) -> pd.DataFrame:
    df = hoja1.copy()
    codcli_col = selected.get("CODCLI")
    rut_col = selected.get("RUT")
    dv_col = selected.get("DV")
    ano_col = selected.get("ANO")
    periodo_col = selected.get("PERIODO")
    estado_col = selected.get("ESTADO_ASIGNATURA")
    course_cols = selected.get("COURSE_FALLBACK_COLS") or []
    df["_CODCLI_KEY"] = df[codcli_col].map(clean_key) if codcli_col else ""
    if rut_col:
        dv_values = df[dv_col].tolist() if dv_col and dv_col in df.columns else [None] * len(df)
        df["_RUT_NORMALIZADO"] = [normalize_rut(num, dv) for num, dv in zip(df[rut_col].tolist(), dv_values)]
    else:
        df["_RUT_NORMALIZADO"] = ""
    df["_ANO_INT"] = numeric_series(df[ano_col]).astype("Int64") if ano_col else pd.Series([pd.NA] * len(df), dtype="Int64")
    df["_PERIODO_INT"] = numeric_series(df[periodo_col]).astype("Int64") if periodo_col else pd.Series([pd.NA] * len(df), dtype="Int64")
    df["_COURSE_KEY"] = ""
    for col in course_cols:
        vals = df[col].map(clean_key)
        mask = (df["_COURSE_KEY"] == "") & (vals != "")
        df.loc[mask, "_COURSE_KEY"] = vals[mask]
    if estado_col:
        df["_ESTADO_ASIG_ORIG"] = df[estado_col].map(lambda x: norm_text(x))
    else:
        df["_ESTADO_ASIG_ORIG"] = ""
    classes = df["_ESTADO_ASIG_ORIG"].map(classify_estado_asignatura)
    df["_ESTADO_ASIG_CLASIFICACION"] = classes.map(lambda x: x[0])
    df["_ESTADO_ASIG_CRITERIO"] = classes.map(lambda x: x[1])
    return df


def build_catalogo_estado_asignatura(detail: pd.DataFrame) -> pd.DataFrame:
    if detail.empty:
        return pd.DataFrame([{
            "valor_original": "",
            "frecuencia_total": 0,
            "frecuencia_2025": 0,
            "clasificacion": "NO_CLASIFICADO",
            "criterio": "Sin detalle Hoja1",
        }])
    total = detail["_ESTADO_ASIG_ORIG"].fillna("").value_counts(dropna=False).to_dict()
    y2025 = detail.loc[detail["_ANO_INT"] == 2025, "_ESTADO_ASIG_ORIG"].fillna("").value_counts(dropna=False).to_dict()
    values = sorted(set(total) | set(y2025))
    rows = []
    for value in values:
        cls, criterio = classify_estado_asignatura(value)
        rows.append({
            "valor_original": value,
            "frecuencia_total": int(total.get(value, 0)),
            "frecuencia_2025": int(y2025.get(value, 0)),
            "clasificacion": cls,
            "criterio": criterio,
        })
    return pd.DataFrame(rows).sort_values(["clasificacion", "frecuencia_2025", "frecuencia_total"], ascending=[True, False, False])


def build_mapping(sheets: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, dict[str, list[str]], dict[str, list[str]], pd.DataFrame]:
    priority = {"Hoja1": 1, "DatosAlumnos": 2, "Matricula_2025": 3}
    evidence_rows = []
    codcli_seen = []
    for sheet, df in sheets.items():
        cols = detect_mapping_cols(df)
        codcli_col, rut_col, dv_col = cols["CODCLI"], cols["RUT"], cols["DV"]
        if not codcli_col:
            continue
        codcli_values = df[codcli_col].map(clean_key)
        if rut_col:
            dv_values = df[dv_col].tolist() if dv_col and dv_col in df.columns else [None] * len(df)
            rut_values = [normalize_rut(num, dv) for num, dv in zip(df[rut_col].tolist(), dv_values)]
        else:
            rut_values = [""] * len(df)
        tmp = pd.DataFrame({"CODCLI": codcli_values, "RUT_NORMALIZADO": rut_values})
        tmp = tmp[tmp["CODCLI"] != ""]
        if not tmp.empty:
            g_codcli = tmp.groupby("CODCLI", dropna=False).agg(
                apariciones=("CODCLI", "size"),
                ruts_validos=("RUT_NORMALIZADO", lambda s: int((s != "").sum())),
            ).reset_index()
            for _, row in g_codcli.iterrows():
                codcli_seen.append({
                    "CODCLI": row["CODCLI"],
                    "hoja_origen": sheet,
                    "apariciones": int(row["apariciones"]),
                    "ruts_validos": int(row["ruts_validos"]),
                })
        tmp = tmp[(tmp["CODCLI"] != "") & (tmp["RUT_NORMALIZADO"] != "")]
        if tmp.empty:
            continue
        grouped = tmp.groupby(["CODCLI", "RUT_NORMALIZADO"], dropna=False).size().reset_index(name="cantidad_apariciones")
        for _, row in grouped.iterrows():
            evidence_rows.append({
                "CODCLI": row["CODCLI"],
                "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
                "hoja_origen": sheet,
                "prioridad_origen": priority.get(sheet, 9),
                "cantidad_apariciones": int(row["cantidad_apariciones"]),
            })
    evidence = pd.DataFrame(evidence_rows)
    codcli_seen_df = pd.DataFrame(codcli_seen)
    if evidence.empty:
        return pd.DataFrame(), {}, {}, codcli_seen_df
    grouped = evidence.groupby(["CODCLI", "RUT_NORMALIZADO"], dropna=False).agg(
        hojas_origen=("hoja_origen", lambda s: safe_join(sorted(set(s)))),
        prioridad_min=("prioridad_origen", "min"),
        cantidad_apariciones=("cantidad_apariciones", "sum"),
    ).reset_index()
    codcli_to_rut_count = grouped.groupby("CODCLI")["RUT_NORMALIZADO"].nunique().to_dict()
    rut_to_codcli_count = grouped.groupby("RUT_NORMALIZADO")["CODCLI"].nunique().to_dict()
    grouped["CODCLI_CON_MULTIPLES_RUT"] = grouped["CODCLI"].map(lambda x: "SI" if codcli_to_rut_count.get(x, 0) > 1 else "NO")
    grouped["RUT_CON_MULTIPLES_CODCLI"] = grouped["RUT_NORMALIZADO"].map(lambda x: "SI" if rut_to_codcli_count.get(x, 0) > 1 else "NO")
    rut_to_codcli: dict[str, list[str]] = defaultdict(list)
    codcli_to_rut: dict[str, list[str]] = defaultdict(list)
    for _, row in grouped.iterrows():
        rut_to_codcli[row["RUT_NORMALIZADO"]].append(row["CODCLI"])
        codcli_to_rut[row["CODCLI"]].append(row["RUT_NORMALIZADO"])
    rut_to_codcli = {k: sorted(set(v)) for k, v in rut_to_codcli.items()}
    codcli_to_rut = {k: sorted(set(v)) for k, v in codcli_to_rut.items()}
    grouped.insert(0, "TIPO_REGISTRO", "MAPEO_CODCLI_RUT")
    return grouped.sort_values(["RUT_NORMALIZADO", "CODCLI"]), rut_to_codcli, codcli_to_rut, codcli_seen_df


def detect_estado_academico_cols(sheet: str, df: pd.DataFrame) -> list[str]:
    cols = []
    for col in df.columns:
        ncol = norm_col(col)
        if sheet == "Hoja1" and ncol in {"ESTADO", "DESCRIPCIONESTADO"}:
            continue
        explicit = ncol in {
            "ESTADOACADEMICO",
            "ESTADOALUMNO",
            "ESTADOMATRICULA",
            "SITUACION",
            "SITUACIONACADEMICA",
            "CONDICION",
            "CONDICIONACADEMICA",
            "MATRICULA",
        }
        contains = any(token in ncol for token in ["ESTADOACADEMICO", "SITUACIONACADEMICA", "CONDICIONACADEMICA", "ESTADOALUMNO", "ESTADOMATRICULA"])
        if explicit or contains:
            cols.append(str(col))
    return cols


def build_estado_academico_mapping(sheets: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    rows = []
    for sheet, df in sheets.items():
        map_cols = detect_mapping_cols(df)
        codcli_col, rut_col, dv_col = map_cols["CODCLI"], map_cols["RUT"], map_cols["DV"]
        estado_cols = detect_estado_academico_cols(sheet, df)
        if not estado_cols or not (codcli_col or rut_col):
            continue
        if codcli_col:
            codcli_values = df[codcli_col].map(clean_key).tolist()
        else:
            codcli_values = [""] * len(df)
        if rut_col:
            dv_values = df[dv_col].tolist() if dv_col and dv_col in df.columns else [None] * len(df)
            rut_values = [normalize_rut(num, dv) for num, dv in zip(df[rut_col].tolist(), dv_values)]
        else:
            rut_values = [""] * len(df)
        for estado_col in estado_cols:
            vals = df[estado_col].map(lambda x: norm_text(x)).tolist()
            textual_col = norm_col(estado_col) != "MATRICULA"
            tmp = pd.DataFrame({
                "CODCLI": codcli_values,
                "RUT_NORMALIZADO": rut_values,
                "VALOR_ESTADO_ACADEMICO": vals,
            })
            tmp = tmp[tmp["VALOR_ESTADO_ACADEMICO"] != ""]
            if tmp.empty:
                continue
            grouped = tmp.groupby(["CODCLI", "RUT_NORMALIZADO", "VALOR_ESTADO_ACADEMICO"], dropna=False).size().reset_index(name="frecuencia")
            for _, row in grouped.iterrows():
                cls, criterio = classify_estado_academico(row["VALOR_ESTADO_ACADEMICO"])
                rows.append({
                    "TIPO_REGISTRO": "ESTADO_ACADEMICO_INSTITUCIONAL",
                    "HOJA_ORIGEN": sheet,
                    "COLUMNA_ORIGEN": estado_col,
                    "ES_CAMPO_TEXTUAL": "SI" if textual_col else "NO",
                    "CODCLI": row["CODCLI"],
                    "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
                    "VALOR_ESTADO_ACADEMICO": row["VALOR_ESTADO_ACADEMICO"],
                    "CLASIFICACION_ESTADO_ACADEMICO": cls,
                    "CRITERIO": criterio,
                    "FRECUENCIA": int(row["frecuencia"]),
                })
    df_estado = pd.DataFrame(rows)
    by_codcli: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_rut: dict[str, list[dict[str, Any]]] = defaultdict(list)
    if not df_estado.empty:
        for _, row in df_estado.iterrows():
            item = row.to_dict()
            if row.get("CODCLI"):
                by_codcli[str(row["CODCLI"])].append(item)
            if row.get("RUT_NORMALIZADO"):
                by_rut[str(row["RUT_NORMALIZADO"])].append(item)
    return df_estado, by_codcli, by_rut


def build_activity_metrics(detail: pd.DataFrame) -> dict[str, dict[str, Any]]:
    metrics: dict[str, dict[str, Any]] = {}
    if detail.empty or "_CODCLI_KEY" not in detail.columns:
        return metrics
    valid = detail[detail["_CODCLI_KEY"] != ""]
    for codcli, group in valid.groupby("_CODCLI_KEY", dropna=False):
        g2025 = group[group["_ANO_INT"] == 2025]
        g2026 = group[group["_ANO_INT"] == 2026]
        gle2025 = group[group["_ANO_INT"].notna() & (group["_ANO_INT"] <= 2025)]
        courses_2025 = set(g2025.loc[g2025["_COURSE_KEY"] != "", "_COURSE_KEY"])
        approved_2025 = set(g2025.loc[(g2025["_COURSE_KEY"] != "") & (g2025["_ESTADO_ASIG_CLASIFICACION"] == "APROBADO"), "_COURSE_KEY"])
        unclass_2025 = set(g2025.loc[(g2025["_COURSE_KEY"] != "") & (g2025["_ESTADO_ASIG_CLASIFICACION"] == "NO_CLASIFICADO"), "_COURSE_KEY"])
        unclass_affects_2025 = unclass_2025 - approved_2025
        periods_2025 = set(int(x) for x in g2025["_PERIODO_INT"].dropna().tolist())
        courses_le2025 = set(gle2025.loc[gle2025["_COURSE_KEY"] != "", "_COURSE_KEY"])
        approved_le2025 = set(gle2025.loc[(gle2025["_COURSE_KEY"] != "") & (gle2025["_ESTADO_ASIG_CLASIFICACION"] == "APROBADO"), "_COURSE_KEY"])
        metrics[str(codcli)] = {
            "registros_2025": int(len(g2025)),
            "registros_2026": int(len(g2026)),
            "periodos_2025": periods_2025,
            "courses_2025": courses_2025,
            "approved_2025": approved_2025,
            "unclass_affects_2025": unclass_affects_2025,
            "courses_le2025": courses_le2025,
            "approved_le2025": approved_le2025,
        }
    return metrics


def union_metrics(codclis: list[str], metrics: dict[str, dict[str, Any]]) -> dict[str, Any]:
    out = {
        "registros_2025": 0,
        "registros_2026": 0,
        "periodos_2025": set(),
        "courses_2025": set(),
        "approved_2025": set(),
        "unclass_affects_2025": set(),
        "courses_le2025": set(),
        "approved_le2025": set(),
    }
    for codcli in codclis:
        item = metrics.get(codcli)
        if not item:
            continue
        out["registros_2025"] += item["registros_2025"]
        out["registros_2026"] += item["registros_2026"]
        for key in ["periodos_2025", "courses_2025", "approved_2025", "unclass_affects_2025", "courses_le2025", "approved_le2025"]:
            out[key] = out[key] | set(item[key])
    return out


def estado_items_for_row(codclis: list[str], rut: str, by_codcli: dict[str, list[dict[str, Any]]], by_rut: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    seen = set()
    items = []
    for codcli in codclis:
        for item in by_codcli.get(codcli, []):
            key = (item.get("HOJA_ORIGEN"), item.get("COLUMNA_ORIGEN"), item.get("CODCLI"), item.get("RUT_NORMALIZADO"), item.get("VALOR_ESTADO_ACADEMICO"))
            if key not in seen:
                seen.add(key)
                items.append(item)
    for item in by_rut.get(rut, []):
        key = (item.get("HOJA_ORIGEN"), item.get("COLUMNA_ORIGEN"), item.get("CODCLI"), item.get("RUT_NORMALIZADO"), item.get("VALOR_ESTADO_ACADEMICO"))
        if key not in seen:
            seen.add(key)
            items.append(item)
    return items


def audit_annual_rows(csv_df: pd.DataFrame, rut_to_codcli: dict[str, list[str]], metrics: dict[str, dict[str, Any]], estado_by_codcli: dict[str, list[dict[str, Any]]], estado_by_rut: dict[str, list[dict[str, Any]]]) -> pd.DataFrame:
    rows = []
    for idx, row in csv_df.iterrows():
        rut = row["RUT_NORMALIZADO"]
        codclis = rut_to_codcli.get(rut, [])
        metric = union_metrics(codclis, metrics)
        estado_items = estado_items_for_row(codclis, rut, estado_by_codcli, estado_by_rut)
        textual_estado = [item for item in estado_items if item.get("ES_CAMPO_TEXTUAL") == "SI"]
        estado_values = sorted(set(item["VALOR_ESTADO_ACADEMICO"] for item in textual_estado if item.get("VALOR_ESTADO_ACADEMICO")))
        estado_identificado = bool(estado_values)
        estado_explicativo = any(is_estado_explicativo(value) for value in estado_values)
        tiene_actividad = metric["registros_2025"] > 0
        calc_c1 = "SI" if 1 in metric["periodos_2025"] else "NO"
        calc_c2 = "SI" if 2 in metric["periodos_2025"] else "NO"
        calc_cursadas = len(metric["courses_2025"])
        calc_aprobadas = len(metric["approved_2025"])
        orig_c1 = normalize_si_no(row["CURSO_1ER_SEM"])
        orig_c2 = normalize_si_no(row["CURSO_2DO_SEM"])
        orig_cursadas = to_int(row["UNIDADES_CURSADAS"])
        orig_aprobadas = to_int(row["UNIDADES_APROBADAS"])
        ok_c1 = orig_c1 == calc_c1
        ok_c2 = orig_c2 == calc_c2
        ok_cursadas = orig_cursadas == calc_cursadas
        ok_aprobadas = orig_aprobadas == calc_aprobadas
        calza = ok_c1 and ok_c2 and ok_cursadas and ok_aprobadas
        no_no_0_0 = orig_c1 == "NO" and orig_c2 == "NO" and orig_cursadas == 0 and orig_aprobadas == 0
        unclass_affects = len(metric["unclass_affects_2025"]) > 0
        if not codclis:
            dictamen = "BLOQUEO_SIN_CODCLI_ASOCIADO"
            obs = "No se logro mapear CODCLI/RUT para buscar evidencia academica."
        elif tiene_actividad and unclass_affects:
            dictamen = "BLOQUEO_ESTADO_ASIGNATURA_NO_CLASIFICADO"
            obs = "Hay actividad 2025 con estado de asignatura no clasificable que puede afectar aprobadas."
        elif tiene_actividad:
            if estado_explicativo and calza:
                dictamen = "OK_CON_ACTIVIDAD_2025_AUNQUE_ESTADO_INACTIVO_O_ELIMINADO"
                obs = "Manda evidencia academica 2025; el estado no anula el calculo."
            elif estado_explicativo and not calza:
                dictamen = "REVISAR_ESTADO_INDICA_SIN_ACTIVIDAD_PERO_EXISTE_ACTIVIDAD_2025"
                obs = "No se fuerza cero: existe evidencia 2025 y el 5809 no calza con el recalculo."
            elif calza:
                dictamen = "OK_CON_ACTIVIDAD_2025"
                obs = "Columnas 16-19 calzan con actividad academica real 2025."
            else:
                dictamen = "DIFERENCIA_CALCULO_2025"
                obs = "Hay evidencia 2025, pero columnas 16-19 no calzan con el recalculo."
        else:
            if estado_explicativo and no_no_0_0:
                dictamen = "OK_SIN_ACTIVIDAD_2025_RESPALDADO_POR_ESTADO"
                obs = "No hay registros 2025 y el estado explica ausencia; se espera NO/NO/0/0."
            elif estado_explicativo and not no_no_0_0:
                dictamen = "REVISAR_SIN_ACTIVIDAD_RESPALDADA_PERO_5809_NO_CALZA"
                obs = "No hay registros 2025 y el estado explica ausencia, pero 5809 no esta en NO/NO/0/0."
            else:
                dictamen = "BLOQUEO_SIN_ESTADO_Y_SIN_EVIDENCIA_2025"
                obs = "No hay actividad 2025 ni estado academico explicativo."
        rows.append({
            "FILA_5809": int(idx + 1),
            "RUT_NORMALIZADO": rut,
            "CODCLI_LISTA": safe_join(codclis, ";"),
            "CODCLI_CANTIDAD": len(codclis),
            "CODCLI_MULTIPLE": "SI" if len(codclis) > 1 else "NO",
            "CODIGO_UNICO": row["CODIGO_UNICO"],
            "PLAN_ESTUDIOS": row["PLAN_ESTUDIOS"],
            "ESTADO_ACADEMICO_IDENTIFICADO": "SI" if estado_identificado else "NO",
            "ESTADO_ACADEMICO_EXPLICATIVO": "SI" if estado_explicativo else "NO",
            "ESTADO_ACADEMICO_RESUMEN": safe_join(estado_values),
            "REGISTROS_2025": metric["registros_2025"],
            "TIENE_ACTIVIDAD_ACADEMICA_2025": "SI" if tiene_actividad else "NO",
            "REGISTROS_2026_MISMO_CODCLI_ALERTA": metric["registros_2026"],
            "CURSO_1ER_SEM_5809": row["CURSO_1ER_SEM"],
            "CURSO_1ER_SEM_CALC": calc_c1,
            "OK_CURSO_1ER_SEM": "SI" if ok_c1 else "NO",
            "CURSO_2DO_SEM_5809": row["CURSO_2DO_SEM"],
            "CURSO_2DO_SEM_CALC": calc_c2,
            "OK_CURSO_2DO_SEM": "SI" if ok_c2 else "NO",
            "UNIDADES_CURSADAS_5809": row["UNIDADES_CURSADAS"],
            "UNIDADES_CURSADAS_2025_CALC": calc_cursadas,
            "DIF_UNIDADES_CURSADAS_2025": None if orig_cursadas is None else orig_cursadas - calc_cursadas,
            "OK_UNIDADES_CURSADAS_2025": "SI" if ok_cursadas else "NO",
            "UNIDADES_APROBADAS_5809": row["UNIDADES_APROBADAS"],
            "UNIDADES_APROBADAS_2025_CALC": calc_aprobadas,
            "DIF_UNIDADES_APROBADAS_2025": None if orig_aprobadas is None else orig_aprobadas - calc_aprobadas,
            "OK_UNIDADES_APROBADAS_2025": "SI" if ok_aprobadas else "NO",
            "CURSOS_2025_DISTINTOS": safe_join(sorted(metric["courses_2025"])),
            "CURSOS_APROBADOS_2025_DISTINTOS": safe_join(sorted(metric["approved_2025"])),
            "ESTADOS_ASIGNATURA_NO_CLASIFICADOS_AFECTAN_APROBADAS": len(metric["unclass_affects_2025"]),
            "CURSOS_ESTADO_NO_CLASIFICADO_AFECTA": safe_join(sorted(metric["unclass_affects_2025"])),
            "5809_NO_NO_0_0": "SI" if no_no_0_0 else "NO",
            "DICTAMEN_ESTADO_Y_AVANCE_2025": dictamen,
            "OBSERVACION_FUNCIONAL": obs,
        })
    return pd.DataFrame(rows)


def audit_accumulated_rows(csv_df: pd.DataFrame, audit_df: pd.DataFrame, rut_to_codcli: dict[str, list[str]], metrics: dict[str, dict[str, Any]]) -> pd.DataFrame:
    rows = []
    annual_by_file = audit_df.set_index("FILA_5809").to_dict("index")
    for idx, row in csv_df.iterrows():
        fila = int(idx + 1)
        rut = row["RUT_NORMALIZADO"]
        codclis = rut_to_codcli.get(rut, [])
        metric = union_metrics(codclis, metrics)
        orig_cursadas_total = to_int(row["UNID_CURSADAS_TOTAL"])
        orig_aprobadas_total = to_int(row["UNID_APROBADAS_TOTAL"])
        calc_cursadas_total = len(metric["courses_le2025"])
        calc_aprobadas_total = len(metric["approved_le2025"])
        if not codclis:
            dictamen = "BLOQUEO"
            motivo = "Sin CODCLI asociado; no es posible auditar acumulado 20-21."
        elif len(codclis) > 1:
            dictamen = "NO_EJECUTABLE_SIN_LLAVE_PLAN_CARRERA"
            motivo = "RUT asociado a multiples CODCLI; se requiere llave plan/carrera para demostrar acumulado sin mezclar trayectorias."
        elif orig_cursadas_total == calc_cursadas_total and orig_aprobadas_total == calc_aprobadas_total:
            dictamen = "OK_ACUMULADO_CALZA"
            motivo = "Acumulado 20-21 calza con cursos distintos hasta cierre 2025 por CODCLI unico."
        elif orig_cursadas_total is None or orig_aprobadas_total is None:
            dictamen = "BLOQUEO"
            motivo = "Valores 20-21 no son numericos evaluables."
        else:
            dictamen = "REVISAR_LOGICA_ACUMULADA"
            motivo = "No calza con conteo acumulado simple por CODCLI unico; mantener separado del dictamen anual."
        rows.append({
            "FILA_5809": fila,
            "RUT_NORMALIZADO": rut,
            "CODCLI_LISTA": annual_by_file.get(fila, {}).get("CODCLI_LISTA", ""),
            "CODCLI_CANTIDAD": len(codclis),
            "CODIGO_UNICO": row["CODIGO_UNICO"],
            "PLAN_ESTUDIOS": row["PLAN_ESTUDIOS"],
            "UNID_CURSADAS_TOTAL_5809": row["UNID_CURSADAS_TOTAL"],
            "UNID_CURSADAS_TOTAL_CALC_LE2025": calc_cursadas_total,
            "DIF_UNID_CURSADAS_TOTAL": None if orig_cursadas_total is None else orig_cursadas_total - calc_cursadas_total,
            "UNID_APROBADAS_TOTAL_5809": row["UNID_APROBADAS_TOTAL"],
            "UNID_APROBADAS_TOTAL_CALC_LE2025": calc_aprobadas_total,
            "DIF_UNID_APROBADAS_TOTAL": None if orig_aprobadas_total is None else orig_aprobadas_total - calc_aprobadas_total,
            "DICTAMEN_ACUMULADO_20_21": dictamen,
            "MOTIVO": motivo,
            "NOTA_SEPARACION": "Este dictamen no bloquea ni reemplaza el dictamen anual 16-19.",
        })
    return pd.DataFrame(rows)


def make_mapping_output(mapping_df: pd.DataFrame, codcli_seen_df: pd.DataFrame, csv_df: pd.DataFrame, audit_df: pd.DataFrame) -> pd.DataFrame:
    frames = []
    if not mapping_df.empty:
        frames.append(mapping_df.copy())
    diag_rows = []
    mapped_codcli = set(mapping_df["CODCLI"]) if not mapping_df.empty and "CODCLI" in mapping_df else set()
    if not codcli_seen_df.empty:
        codcli_summary = codcli_seen_df.groupby("CODCLI", dropna=False).agg(
            hojas_origen=("hoja_origen", lambda s: safe_join(sorted(set(s)))),
            apariciones=("apariciones", "sum"),
            ruts_validos=("ruts_validos", "sum"),
        ).reset_index()
        for _, row in codcli_summary[codcli_summary["CODCLI"].ne("") & ~codcli_summary["CODCLI"].isin(mapped_codcli)].iterrows():
            diag_rows.append({
                "TIPO_REGISTRO": "DIAGNOSTICO_CODCLI_SIN_RUT",
                "CODCLI": row["CODCLI"],
                "RUT_NORMALIZADO": "",
                "hojas_origen": row["hojas_origen"],
                "cantidad_apariciones": int(row["apariciones"]),
                "DETALLE": "CODCLI observado sin RUT normalizable en fuentes PROMEDIOS.",
            })
    no_codcli = audit_df[audit_df["CODCLI_CANTIDAD"] == 0][["FILA_5809", "RUT_NORMALIZADO", "CODIGO_UNICO", "PLAN_ESTUDIOS"]].copy()
    for _, row in no_codcli.iterrows():
        diag_rows.append({
            "TIPO_REGISTRO": "DIAGNOSTICO_RUT_5809_SIN_CODCLI_ASOCIADO",
            "CODCLI": "",
            "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
            "FILA_5809": row["FILA_5809"],
            "CODIGO_UNICO": row["CODIGO_UNICO"],
            "PLAN_ESTUDIOS": row["PLAN_ESTUDIOS"],
            "DETALLE": "RUT del 5809 sin CODCLI asociado.",
        })
    for _, row in audit_df[audit_df["TIENE_ACTIVIDAD_ACADEMICA_2025"] == "SI"][["CODCLI_LISTA", "RUT_NORMALIZADO"]].iterrows():
        for codcli in str(row["CODCLI_LISTA"]).split(";"):
            if codcli:
                diag_rows.append({
                    "TIPO_REGISTRO": "DIAGNOSTICO_CODCLI_CON_ACTIVIDAD_2025",
                    "CODCLI": codcli,
                    "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
                    "DETALLE": "CODCLI con registros academicos ANO=2025.",
                })
    for _, row in audit_df[(audit_df["TIENE_ACTIVIDAD_ACADEMICA_2025"] == "NO") & (audit_df["CODCLI_CANTIDAD"] > 0)][["CODCLI_LISTA", "RUT_NORMALIZADO"]].iterrows():
        for codcli in str(row["CODCLI_LISTA"]).split(";"):
            if codcli:
                diag_rows.append({
                    "TIPO_REGISTRO": "DIAGNOSTICO_CODCLI_SIN_ACTIVIDAD_2025",
                    "CODCLI": codcli,
                    "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
                    "DETALLE": "CODCLI sin registros academicos ANO=2025.",
                })
    for _, row in audit_df[(audit_df["REGISTROS_2026_MISMO_CODCLI_ALERTA"] > 0) & (audit_df["CODCLI_CANTIDAD"] > 0)][["CODCLI_LISTA", "RUT_NORMALIZADO", "REGISTROS_2026_MISMO_CODCLI_ALERTA"]].iterrows():
        for codcli in str(row["CODCLI_LISTA"]).split(";"):
            if codcli:
                diag_rows.append({
                    "TIPO_REGISTRO": "DIAGNOSTICO_CODCLI_CON_ACTIVIDAD_2026_ALERTA",
                    "CODCLI": codcli,
                    "RUT_NORMALIZADO": row["RUT_NORMALIZADO"],
                    "REGISTROS_2026_MISMO_CODCLI_ALERTA": row["REGISTROS_2026_MISMO_CODCLI_ALERTA"],
                    "DETALLE": "Actividad 2026 detectada solo como alerta; no se usa para bloquear ni calcular 2025.",
                })
    if diag_rows:
        frames.append(pd.DataFrame(diag_rows))
    if frames:
        return pd.concat(frames, ignore_index=True, sort=False)
    return pd.DataFrame([{"TIPO_REGISTRO": "SIN_MAPEO", "DETALLE": "No se construyo mapeo CODCLI/RUT."}])


def build_regla_funcional() -> pd.DataFrame:
    rows = [
        ("A_REGLA_OFICIAL", "El instructivo exige informar actividad anual 2025 y acumulado hasta cierre 2025."),
        ("A_REGLA_OFICIAL", "La vigencia del proceso no debe confundirse con estado academico institucional."),
        ("A_REGLA_OFICIAL", "Si un estudiante tuvo actividad academica en 2025, debe informarse su avance."),
        ("B_DATO_OBSERVADO", "Se observa estado academico desde PROMEDIOSDEALUMNOS o fuentes institucionales, registros academicos 2025 y valores actuales del 5809."),
        ("C_IMPLEMENTACION_TECNICA", "Primero se identifica RUT/CODCLI; despues se busca evidencia academica 2025; columnas 16-19 se recalculan solo con ANO=2025."),
        ("C_IMPLEMENTACION_TECNICA", "UNIDADES_CURSADAS cuenta CODRAMO/asignaturas distintas 2025; UNIDADES_APROBADAS cuenta distintas con estado de asignatura APROBADO por patron explicito."),
        ("D_DECISION_INTERNA_GOBERNADA", "Estado ELIMINADO o SIN ACTIVIDAD ACADEMICA explica ausencia si no hay registros 2025; no reemplaza evidencia academica."),
        ("D_DECISION_INTERNA_GOBERNADA", "Esta decision interna no se presenta como regla oficial y no fuerza ceros por estado academico."),
        ("E_PENDIENTE", "Casos sin estado y sin registros 2025, diferencias de calculo, estados de asignatura no clasificables y acumulado 20-21 pendiente."),
        ("CASO_A", "Con actividad academica 2025 manda evidencia academica; el estado no bloquea ni anula el calculo."),
        ("CASO_B", "Sin actividad 2025 y estado explicativo: OK si 5809 esta en NO/NO/0/0."),
        ("CASO_C", "Estado indica sin actividad/eliminado pero hay registros 2025: no forzar cero; recalcular con evidencia."),
        ("CASO_D", "Sin estado academico explicativo y sin evidencia 2025: bloqueo real anual."),
        ("CASO_E", "Hay evidencia 2025 pero calculo no calza: diferencia de calculo 2025."),
    ]
    return pd.DataFrame(rows, columns=["BLOQUE", "REGLA"])


def build_validaciones(inputs: dict[str, Any], selected: dict[str, Any], csv_df: pd.DataFrame, sheets: dict[str, pd.DataFrame], audit_df: pd.DataFrame, accum_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    rows.append({"VALIDACION": "CSV_5809_LECTURA", "RESULTADO": "OK" if len(csv_df.columns) >= 23 else "BLOQUEO", "DETALLE": f"Leido sin encabezado, separador punto y coma, cp1252. Filas={len(csv_df)}, columnas_originales=22."})
    rows.append({"VALIDACION": "PROMEDIOS_TODAS_LAS_HOJAS", "RESULTADO": "OK" if sheets else "BLOQUEO", "DETALLE": f"Hojas leidas: {safe_join(list(sheets.keys()))}."})
    required = ["CODCLI", "ANO", "PERIODO", "CODRAMO_ASIGNATURA", "ESTADO_ASIGNATURA"]
    missing = [key for key in required if not selected.get(key)]
    rows.append({"VALIDACION": "HOJA1_DETALLE_ACADEMICO", "RESULTADO": "OK" if not missing else "BLOQUEO", "DETALLE": "Columnas seleccionadas: " + json.dumps({k: selected.get(k) for k in required}, ensure_ascii=False)})
    estado_col = selected.get("ESTADO_ASIGNATURA") or ""
    rows.append({"VALIDACION": "NO_CONFUNDIR_ESTADO_ASIGNATURA_CON_ESTADO_ACADEMICO", "RESULTADO": "OK" if "ACADEMIC" not in norm_col(estado_col) else "BLOQUEO", "DETALLE": f"Columna usada para estado de asignatura: {estado_col}."})
    rows.append({"VALIDACION": "EXCLUSION_2026_DEMOSTRADA", "RESULTADO": "OK" if selected.get("ANO") else "BLOQUEO", "DETALLE": "El calculo anual usa solo ANO=2025. Registros 2026 se exponen como alerta separada."})
    rows.append({"VALIDACION": "ESTADO_NO_FUERZA_CEROS", "RESULTADO": "OK", "DETALLE": "Si hay registros academicos 2025, el dictamen usa evidencia academica aunque el estado sea eliminado o sin actividad."})
    rows.append({"VALIDACION": "ACUMULADO_SEPARADO", "RESULTADO": "OK", "DETALLE": f"Hoja separada 13_ACUMULADO_20_21_SEPARADO. Pendientes={int((accum_df['DICTAMEN_ACUMULADO_20_21'] != 'OK_ACUMULADO_CALZA').sum())}."})
    rows.append({"VALIDACION": "CSV_5809_NO_MODIFICADO", "RESULTADO": "OK", "DETALLE": "Este script solo lee el CSV 5809 y genera auditoria; no escribe ni corrige la fuente."})
    rows.append({"VALIDACION": "RUT_NORMALIZADO", "RESULTADO": "OK", "DETALLE": "RUT 5809 normalizado como NUM_DOCUMENTO-DV."})
    rows.append({"VALIDACION": "DICTAMENES_ANUALES_GENERADOS", "RESULTADO": "OK", "DETALLE": f"Filas auditadas={len(audit_df)}."})
    if inputs.get("instructivo_usado"):
        rows.append({"VALIDACION": "INSTRUCTIVO_LOCALIZADO", "RESULTADO": "OK", "DETALLE": str(inputs["instructivo_usado"])})
    else:
        rows.append({"VALIDACION": "INSTRUCTIVO_LOCALIZADO", "RESULTADO": "ADVERTENCIA", "DETALLE": "No se encontro el instructivo por ruta exacta ni busqueda local; se registra sin detener la auditoria tecnica."})
    return pd.DataFrame(rows)


def build_global(audit_df: pd.DataFrame, accum_df: pd.DataFrame, selected: dict[str, Any]) -> tuple[pd.DataFrame, dict[str, Any]]:
    ok_dictamens = {
        "OK_CON_ACTIVIDAD_2025",
        "OK_CON_ACTIVIDAD_2025_AUNQUE_ESTADO_INACTIVO_O_ELIMINADO",
        "OK_SIN_ACTIVIDAD_2025_RESPALDADO_POR_ESTADO",
    }
    diff_dictamens = {
        "DIFERENCIA_CALCULO_2025",
        "REVISAR_ESTADO_INDICA_SIN_ACTIVIDAD_PERO_EXISTE_ACTIVIDAD_2025",
        "REVISAR_SIN_ACTIVIDAD_RESPALDADA_PERO_5809_NO_CALZA",
    }
    total = len(audit_df)
    filas_sin_codcli = int((audit_df["CODCLI_CANTIDAD"] == 0).sum())
    filas_con_codcli = total - filas_sin_codcli
    filas_con_estado = int((audit_df["ESTADO_ACADEMICO_IDENTIFICADO"] == "SI").sum())
    filas_con_actividad = int((audit_df["TIENE_ACTIVIDAD_ACADEMICA_2025"] == "SI").sum())
    filas_sin_act_resp = int((audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"] == "OK_SIN_ACTIVIDAD_2025_RESPALDADO_POR_ESTADO").sum())
    filas_act_estado_inactivo = int(((audit_df["TIENE_ACTIVIDAD_ACADEMICA_2025"] == "SI") & (audit_df["ESTADO_ACADEMICO_EXPLICATIVO"] == "SI")).sum())
    filas_ok_anual = int(audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].isin(ok_dictamens).sum())
    filas_diferencias = int(audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].isin(diff_dictamens).sum())
    filas_bloqueo_real = int(audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].str.startswith("BLOQUEO").sum())
    filas_acum_pendiente = int((accum_df["DICTAMEN_ACUMULADO_20_21"] != "OK_ACUMULADO_CALZA").sum())
    estados_asig_no_clasif = int((audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"] == "BLOQUEO_ESTADO_ASIGNATURA_NO_CLASIFICADO").sum())
    exclusion_2026_demostrada = bool(selected.get("ANO"))
    anual_listo = filas_bloqueo_real == 0 and filas_diferencias == 0 and estados_asig_no_clasif == 0 and filas_sin_codcli == 0 and exclusion_2026_demostrada
    acumulado_listo = filas_acum_pendiente == 0
    dictamen_anual = "LISTO_ANUAL_16_19_CON_REGLA_CORREGIDA" if anual_listo else "NO_LISTO_ANUAL_16_19_CON_PENDIENTES"
    dictamen_acumulado = "ACUMULADO_20_21_DEMOSTRADO" if acumulado_listo else "ACUMULADO_20_21_PENDIENTE_O_NO_EJECUTABLE"
    declaracion = "LISTO_PARA_CARGA" if anual_listo and acumulado_listo else "NO_LISTO_PARA_CARGA"
    data = {
        "TOTAL_FILAS_5809": total,
        "FILAS_CON_CODCLI_ASOCIADO": filas_con_codcli,
        "FILAS_SIN_CODCLI": filas_sin_codcli,
        "FILAS_CON_ESTADO_ACADEMICO_IDENTIFICADO": filas_con_estado,
        "FILAS_CON_ACTIVIDAD_ACADEMICA_2025": filas_con_actividad,
        "FILAS_SIN_ACTIVIDAD_2025_RESPALDADAS_POR_ESTADO": filas_sin_act_resp,
        "FILAS_CON_ACTIVIDAD_2025_AUNQUE_ESTADO_ELIMINADO_O_SIN_ACTIVIDAD": filas_act_estado_inactivo,
        "FILAS_OK_ANUAL_16_19": filas_ok_anual,
        "FILAS_CON_DIFERENCIAS_16_19": filas_diferencias,
        "FILAS_CON_BLOQUEO_REAL_ANUAL": filas_bloqueo_real,
        "FILAS_CON_ACUMULADO_PENDIENTE": filas_acum_pendiente,
        "FILAS_ESTADO_ASIGNATURA_NO_CLASIFICADO_AFECTA_APROBADAS": estados_asig_no_clasif,
        "EXCLUSION_2026_DEMOSTRADA": "SI" if exclusion_2026_demostrada else "NO",
        "DICTAMEN_ANUAL_2025": dictamen_anual,
        "DICTAMEN_ACUMULADO_20_21": dictamen_acumulado,
        "DECLARACION_CARGA": declaracion,
    }
    rows = [{"INDICADOR": key, "VALOR": value} for key, value in data.items()]
    rows.insert(0, {"INDICADOR": "SUBPROYECTO", "VALOR": SUBPROYECTO})
    rows.insert(0, {"INDICADOR": "PROCESO", "VALOR": PROCESO})
    rows.insert(2, {"INDICADOR": "ANIO_REFERENCIA_DATOS", "VALOR": ANIO_REFERENCIA})
    rows.append({"INDICADOR": "DECLARACION_FUNCIONAL", "VALOR": "El estado academico solo explica ausencia real de registros 2025; no calcula ni fuerza unidades cursadas/aprobadas."})
    return pd.DataFrame(rows), data


def find_instructivo() -> Path | None:
    direct = [
        REPO_ROOT / "Instructivo_Avance Curricular SIES - 2026.txt",
        REPO_ROOT / "Instructivo Avance Curricular SIES - 2026.txt",
    ]
    for path in direct:
        if path.exists():
            return path
    matches = sorted(REPO_ROOT.rglob("*Instructivo*Avance*Curricular*SIES*2026*.txt"))
    return matches[0] if matches else None


def build_fuentes(instructivo: Path | None, desktop_dir: Path, script_path: Path) -> pd.DataFrame:
    fuentes = [
        ("FUENTE_OFICIAL_INSTRUCTIVO", instructivo),
        ("FUENTE_TECNICA_AUDITORIA_CORREGIDA", AUDITORIA_CORREGIDA),
        ("FUENTE_TECNICA_MAPEO_RESOLUCION", MAPEO_RESOLUCION),
        ("FUENTE_PROMEDIOSDEALUMNOS_ACTUALIZADO", PROMEDIOS),
        ("FUENTE_CSV_5809_TECNICO", CSV_5809),
        ("SCRIPT_VALIDACION", script_path),
        ("SALIDA_EXCEL", EXCEL_OUT),
        ("SALIDA_MARKDOWN", MD_OUT),
        ("SALIDA_MANIFEST", MANIFEST_OUT),
        ("CARPETA_ESCRITORIO", desktop_dir),
    ]
    rows = []
    for rol, path in fuentes:
        rows.append({
            "ROL": rol,
            "RUTA": str(path) if path else "",
            "EXISTE": "SI" if path and path.exists() else "NO",
            "SHA256": sha256_file(path) if path and path.is_file() else "",
            "OBSERVACION": "Solo lectura para fuentes; salidas generadas por este script.",
        })
    return pd.DataFrame(rows)


def write_markdown(path: Path, global_data: dict[str, Any], resumen_dictamen: pd.DataFrame, fuentes_df: pd.DataFrame) -> None:
    lines = [
        "# Informe validacion estado y avance 2025 - 5809",
        "",
        f"- Proceso: {PROCESO}",
        f"- Subproyecto: {SUBPROYECTO}",
        f"- Ano de referencia: {ANIO_REFERENCIA}",
        "",
        "## Dictamen global",
    ]
    for key, value in global_data.items():
        lines.append(f"- {key}: {value}")
    lines.extend([
        "",
        "## Regla funcional corregida",
        "",
        "El estado academico institucional se usa como contexto y ruta de busqueda. No se usa para calcular ni forzar unidades cursadas/aprobadas.",
        "Si existe actividad academica 2025, manda la evidencia academica. Si no existe actividad academica 2025, un estado ELIMINADO o SIN ACTIVIDAD ACADEMICA puede respaldar la ausencia.",
        "Las columnas 16-19 se dictaminan separadas del acumulado 20-21.",
        "",
        "## Resumen dictamen anual",
        "",
    ])
    for _, row in resumen_dictamen.iterrows():
        lines.append(f"- {row['DICTAMEN_ESTADO_Y_AVANCE_2025']}: {row['FILAS']}")
    lines.extend([
        "",
        "## Fuentes",
        "",
    ])
    for _, row in fuentes_df.iterrows():
        lines.append(f"- {row['ROL']}: {row['RUTA']} ({row['EXISTE']})")
    lines.extend([
        "",
        "## Declaracion",
        "",
        "No se corrigio el CSV 5809. No se modificaron fuentes originales. No se uso Internet. La decision interna gobernada no se presenta como regla oficial.",
    ])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def sanitize_excel_value(value: Any) -> Any:
    if isinstance(value, str):
        return re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]", "", value)
    return value


def sanitize_df(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame([{"SIN_FILAS": "Sin registros para esta hoja"}])
    out = df.copy()
    for col in out.columns:
        if out[col].dtype == object:
            out[col] = out[col].map(sanitize_excel_value)
    return out


def format_excel(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")
    bad_fill = PatternFill("solid", fgColor="FFC7CE")
    header_font = Font(color="FFFFFF", bold=True)
    ok_font = Font(color="006100", bold=True)
    warn_font = Font(color="9C6500", bold=True)
    bad_font = Font(color="9C0006", bold=True)
    border = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))
    align = Alignment(vertical="top", wrap_text=True)
    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column
        ws.freeze_panes = "A2"
        if max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align
                cell.border = border
            ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, max_col + 1):
            width = 12
            for row_idx in range(1, min(max_row, 300) + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None:
                    width = max(width, len(str(value)) + 2)
            header = str(ws.cell(1, col_idx).value or "").upper()
            if any(token in header for token in ["RUTA", "OBSERVACION", "MOTIVO", "REGLA", "DETALLE", "RESUMEN", "CURSOS"]):
                width = max(width, 45)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(width, 70)
        for row_idx in range(2, min(max_row, 3000) + 1):
            row_text = " ".join(str(ws.cell(row_idx, col_idx).value or "").upper() for col_idx in range(1, max_col + 1))
            fill = font = None
            if "BLOQUEO" in row_text or "NO_LISTO" in row_text:
                fill, font = bad_fill, bad_font
            elif "REVISAR" in row_text or "DIFERENCIA" in row_text or "PENDIENTE" in row_text or "NO_EJECUTABLE" in row_text:
                fill, font = warn_fill, warn_font
            elif "OK" in row_text or "LISTO" in row_text:
                fill, font = ok_fill, ok_font
            if fill:
                for col_idx in range(1, max_col + 1):
                    ws.cell(row_idx, col_idx).fill = fill
                    ws.cell(row_idx, col_idx).font = font
            for col_idx in range(1, max_col + 1):
                ws.cell(row_idx, col_idx).alignment = align
                ws.cell(row_idx, col_idx).border = border
    wb.save(path)


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    tmp_path = path.with_name(path.stem + ".tmp.xlsx")
    for candidate in [tmp_path, path]:
        if candidate.exists():
            candidate.unlink()
    with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            sanitize_df(df).to_excel(writer, sheet_name=sheet_name, index=False)
    format_excel(tmp_path)
    wb = load_workbook(tmp_path, read_only=True)
    wb.close()
    tmp_path.replace(path)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    desktop_dir = Path.home() / f"Desktop/AVANCE_CURRICULAR_2026_VALIDACION_ESTADO_Y_AVANCE_2025_5809_{timestamp}"
    desktop_dir.mkdir(parents=True, exist_ok=True)
    script_path = Path(__file__).resolve()
    missing_critical = [str(path) for path in [PROMEDIOS, CSV_5809] if not path.exists()]
    if missing_critical:
        raise FileNotFoundError("Faltan fuentes criticas: " + "; ".join(missing_critical))
    instructivo = find_instructivo()
    raw_csv = pd.read_csv(CSV_5809, sep=";", header=None, dtype=str, encoding="cp1252", keep_default_na=False)
    if raw_csv.shape[1] != 22:
        raise ValueError(f"CSV 5809 debe tener 22 columnas; detectadas {raw_csv.shape[1]}")
    csv_df = raw_csv.copy()
    csv_df.columns = CSV_COLUMNS
    csv_df["RUT_NORMALIZADO"] = [normalize_rut(num, dv) for num, dv in zip(csv_df["NUM_DOCUMENTO"], csv_df["DV"])]
    xls = pd.ExcelFile(PROMEDIOS, engine="openpyxl")
    sheets = {
        sheet: pd.read_excel(PROMEDIOS, sheet_name=sheet, dtype=str, keep_default_na=False, engine="openpyxl")
        for sheet in xls.sheet_names
    }
    hoja1 = sheets.get("Hoja1", pd.DataFrame())
    selected, estado_eval = detect_hoja1_columns(hoja1)
    detail = prepare_hoja1(hoja1, selected)
    catalogo_estado = build_catalogo_estado_asignatura(detail)
    mapping_df, rut_to_codcli, codcli_to_rut, codcli_seen_df = build_mapping(sheets)
    estado_map_df, estado_by_codcli, estado_by_rut = build_estado_academico_mapping(sheets)
    metrics = build_activity_metrics(detail)
    audit_df = audit_annual_rows(csv_df, rut_to_codcli, metrics, estado_by_codcli, estado_by_rut)
    accum_df = audit_accumulated_rows(csv_df, audit_df, rut_to_codcli, metrics)
    mapping_out = make_mapping_output(mapping_df, codcli_seen_df, csv_df, audit_df)
    global_df, global_data = build_global(audit_df, accum_df, selected)
    resumen_dictamen = audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].value_counts(dropna=False).rename_axis("DICTAMEN_ESTADO_Y_AVANCE_2025").reset_index(name="FILAS")
    resumen_dictamen = pd.concat([resumen_dictamen, pd.DataFrame([{"DICTAMEN_ESTADO_Y_AVANCE_2025": "TOTAL", "FILAS": len(audit_df)}])], ignore_index=True)
    if estado_map_df.empty:
        resumen_estado = pd.DataFrame([{"VALOR_ESTADO_ACADEMICO": "SIN_ESTADOS_DETECTADOS", "FILAS": 0}])
    else:
        resumen_estado = estado_map_df.groupby(["VALOR_ESTADO_ACADEMICO", "CLASIFICACION_ESTADO_ACADEMICO", "CRITERIO"], dropna=False).agg(
            FRECUENCIA=("FRECUENCIA", "sum"),
            CODCLI_DISTINTOS=("CODCLI", lambda s: int(pd.Series(s).replace("", pd.NA).dropna().nunique())),
            RUT_DISTINTOS=("RUT_NORMALIZADO", lambda s: int(pd.Series(s).replace("", pd.NA).dropna().nunique())),
        ).reset_index().sort_values(["CLASIFICACION_ESTADO_ACADEMICO", "FRECUENCIA"], ascending=[True, False])
    fuentes_df = build_fuentes(instructivo, desktop_dir, script_path)
    validaciones_df = build_validaciones({"instructivo_usado": instructivo}, selected, csv_df, sheets, audit_df, accum_df)
    columnas_df = build_column_diagnostics(sheets, selected, estado_eval)
    sheets_out = {
        "00_DICTAMEN_GLOBAL": global_df,
        "01_VALIDACIONES": validaciones_df,
        "02_REGLA_FUNCIONAL_CORREGIDA": build_regla_funcional(),
        "03_COLUMNAS_DETECTADAS": columnas_df,
        "04_CATALOGO_ESTADO_ASIGNATURA": catalogo_estado,
        "05_MAPEO_ESTADO_ACADEMICO": estado_map_df,
        "06_MAPEO_CODCLI_RUT": mapping_out,
        "07_AUDITORIA_ESTADO_Y_AVANCE_2025": audit_df,
        "08_OK_CON_ACTIVIDAD_2025": audit_df[audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].isin(["OK_CON_ACTIVIDAD_2025", "OK_CON_ACTIVIDAD_2025_AUNQUE_ESTADO_INACTIVO_O_ELIMINADO"])],
        "09_OK_SIN_ACTIVIDAD_RESPALDADA": audit_df[audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"] == "OK_SIN_ACTIVIDAD_2025_RESPALDADO_POR_ESTADO"],
        "10_REVISAR_CON_ACTIVIDAD_Y_ESTADO_INACTIVO": audit_df[audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"] == "REVISAR_ESTADO_INDICA_SIN_ACTIVIDAD_PERO_EXISTE_ACTIVIDAD_2025"],
        "11_DIFERENCIAS_CALCULO_2025": audit_df[audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].isin(["DIFERENCIA_CALCULO_2025", "REVISAR_SIN_ACTIVIDAD_RESPALDADA_PERO_5809_NO_CALZA"])],
        "12_BLOQUEOS_REALES_ANUAL": audit_df[audit_df["DICTAMEN_ESTADO_Y_AVANCE_2025"].str.startswith("BLOQUEO")],
        "13_ACUMULADO_20_21_SEPARADO": accum_df,
        "14_RESUMEN_DICTAMEN_ANUAL": resumen_dictamen,
        "15_RESUMEN_ESTADO_ACADEMICO": resumen_estado,
        "FUENTES": fuentes_df,
    }
    write_excel(EXCEL_OUT, sheets_out)
    fuentes_df = build_fuentes(instructivo, desktop_dir, script_path)
    sheets_out["FUENTES"] = fuentes_df
    write_markdown(MD_OUT, global_data, resumen_dictamen, fuentes_df)
    manifest = {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": timestamp,
        "anio_referencia_datos": ANIO_REFERENCIA,
        "salidas": {
            "excel": str(EXCEL_OUT),
            "markdown": str(MD_OUT),
            "manifest": str(MANIFEST_OUT),
            "script": str(script_path),
            "desktop_dir": str(desktop_dir),
        },
        "fuentes": fuentes_df.to_dict("records"),
        "dictamen_global": global_data,
        "hojas_excel": list(sheets_out.keys()),
        "regla_implementada": "Estado academico solo explica ausencia de registros 2025; no calcula ni fuerza unidades 16-19.",
        "csv_5809_modificado": False,
        "fuentes_originales_modificadas": False,
        "internet_usado": False,
    }
    MANIFEST_OUT.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=json_safe) + "\n", encoding="utf-8")
    copied = []
    for path in [EXCEL_OUT, MD_OUT, MANIFEST_OUT, script_path]:
        target = desktop_dir / path.name
        shutil.copy2(path, target)
        copied.append(str(target))
    print("VALIDACION_ESTADO_Y_AVANCE_2025_5809_GENERADA")
    print(f"Excel: {EXCEL_OUT}")
    print(f"Markdown: {MD_OUT}")
    print(f"Manifest: {MANIFEST_OUT}")
    print(f"Escritorio: {desktop_dir}")
    print(json.dumps(global_data, ensure_ascii=False, indent=2, default=json_safe))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
