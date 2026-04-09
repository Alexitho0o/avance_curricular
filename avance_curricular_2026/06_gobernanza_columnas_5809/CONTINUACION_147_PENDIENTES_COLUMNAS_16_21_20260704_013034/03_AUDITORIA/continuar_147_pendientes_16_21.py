#!/usr/bin/env python3
"""Continuacion controlada de los 147 pendientes de columnas 16-21.

No genera CSV de carga ni SIES_READY. Produce una matriz de trabajo auditable
solo si las validaciones de conservacion pasan.
"""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import os
import re
import shutil
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook


TZ = ZoneInfo("America/Santiago")
BASE = Path("/Users/alexi/Documents/GitHub/avance_curricular")
OUT_ROOT = BASE / "avance_curricular_2026" / "06_gobernanza_columnas_5809"
STAMP = datetime.now(TZ).strftime("%Y%m%d_%H%M%S")
RUN_DIR = OUT_ROOT / f"CONTINUACION_147_PENDIENTES_COLUMNAS_16_21_{STAMP}"
RESULTS = RUN_DIR / "02_RESULTADOS"
AUDIT = RUN_DIR / "03_AUDITORIA"
REPORTS = RUN_DIR / "04_REPORTES"

MATRIX_IN = (
    OUT_ROOT
    / "MATRIZ_FINAL_TRABAJO_2224_CALCULADOS_147_PENDIENTES_NO_CARGA_20260704_011030"
    / "02_RESULTADOS"
    / "MATRIZ_FINAL_TRABAJO_2224_CALCULADOS_147_PENDIENTES_NO_CARGA.xlsx"
)
FUENTES = BASE / "avance_curricular_2026" / "00_fuentes_congeladas" / "CARGA_CONGELADA_20260626_005826" / "originales"
PRECARGA_5809 = FUENTES / "5809_Precarga Matrícula Avance Curricular 2026.csv"
PRECARGA_5810 = FUENTES / "5810_Precarga Carreras Avance Curricular 20268.csv"
PROMEDIOS = FUENTES / "PROMEDIOSDEALUMNOS_7804.xlsx"
INSTRUCTIVO = FUENTES / "Instructivo_Avance Curricular SIES - 2026.txt"
CODIGO_GOB = BASE / "codigo_gobernanza_v2.py"

TARGET_FIELDS = [
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
]
PROP_FIELDS = [f"{c}_PROPUESTO" for c in TARGET_FIELDS]

KEYWORDS = [
    "promedios",
    "alumnos",
    "avance",
    "historico",
    "histórico",
    "historial",
    "asignaturas",
    "asignatura",
    "ramos",
    "ramo",
    "notas",
    "malla",
    "matricula",
    "matrícula",
    "matriz",
    "carga",
    "congelada",
]
SUFFIXES = {".xlsx", ".xlsm", ".csv", ".tsv", ".txt", ".parquet"}


def now_iso() -> str:
    return datetime.now(TZ).isoformat(timespec="seconds")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def norm_text(value: Any) -> str:
    text = str(value or "").strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Z0-9]+", "", text)


def norm_rut_body_dv(body: Any, dv: Any = "") -> str:
    body_s = re.sub(r"[^0-9Kk]", "", str(body or "").upper())
    dv_s = re.sub(r"[^0-9Kk]", "", str(dv or "").upper())
    if dv_s and body_s.endswith(dv_s) and len(body_s) > len(dv_s):
        return body_s
    return body_s + dv_s


def split_rut(value: Any) -> str:
    return re.sub(r"[^0-9Kk]", "", str(value or "").upper())


def write_csv(path: Path, df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False, encoding="utf-8", lineterminator="\n")


def write_tsv(path: Path, df: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, sep="\t", index=False, encoding="utf-8", lineterminator="\n")


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def read_csv_flexible(path: Path, nrows: int | None = None) -> pd.DataFrame:
    encodings = ["utf-8-sig", "utf-8", "latin-1"]
    seps = [None, ";", ",", "\t", "|"]
    last: Exception | None = None
    for enc in encodings:
        for sep in seps:
            try:
                return pd.read_csv(
                    path,
                    sep=sep,
                    engine="python" if sep is None else "c",
                    encoding=enc,
                    dtype=str,
                    keep_default_na=False,
                    nrows=nrows,
                    on_bad_lines="skip",
                )
            except Exception as exc:  # noqa: PERF203
                last = exc
    raise last or ValueError(f"No se pudo leer {path}")


def load_function():
    if str(BASE) not in sys.path:
        sys.path.insert(0, str(BASE))
    spec = importlib.util.spec_from_file_location("codigo_gobernanza_v2", CODIGO_GOB)
    if spec is None or spec.loader is None:
        raise ImportError(f"No se pudo importar {CODIGO_GOB}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module.construir_resumen_historico


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            header = str(col_cells[0].value or "")
            max_len = min(max([len(str(c.value or "")) for c in col_cells[:200]] + [len(header)]) + 2, 70)
            ws.column_dimensions[col_cells[0].column_letter].width = max(10, max_len)
    wb.save(path)


def write_xlsx(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(path)


def required_files() -> list[Path]:
    return [CODIGO_GOB, MATRIX_IN, PRECARGA_5809, PRECARGA_5810, PROMEDIOS, INSTRUCTIVO]


def load_5810() -> pd.DataFrame:
    return pd.read_csv(PRECARGA_5810, sep=";", encoding="latin-1", dtype=str, keep_default_na=False)


def load_matrix() -> pd.DataFrame:
    return pd.read_excel(MATRIX_IN, sheet_name="MATRIZ_TRABAJO_2371", dtype=str, keep_default_na=False)


def validate_initial_matrix(df: pd.DataFrame) -> list[dict[str, str]]:
    validations: list[dict[str, str]] = []

    def add(name: str, ok: bool, detail: str) -> None:
        validations.append({"VALIDACION": name, "RESULTADO": "OK" if ok else "ERROR", "DETALLE": detail})

    add("TOTAL_FILAS_2371", len(df) == 2371, str(len(df)))
    add("EXISTE_ESTADO_PREPARACION_16_21", "ESTADO_PREPARACION_16_21" in df.columns, ",".join(df.columns))
    add("CALCULADO_NO_CARGA_2224", int((df.get("ESTADO_PREPARACION_16_21", "") == "CALCULADO_NO_CARGA").sum()) == 2224, str((df.get("ESTADO_PREPARACION_16_21", "") == "CALCULADO_NO_CARGA").sum()))
    add("PENDIENTE_NO_CARGA_147", int((df.get("ESTADO_PREPARACION_16_21", "") == "PENDIENTE_NO_CARGA").sum()) == 147, str((df.get("ESTADO_PREPARACION_16_21", "") == "PENDIENTE_NO_CARGA").sum()))
    pend = df[df.get("ESTADO_PREPARACION_16_21", "") == "PENDIENTE_NO_CARGA"]
    counts = pend.get("ESTADO_CALCULO", pd.Series(dtype=str)).value_counts().to_dict()
    expected = {
        "BLOQUEADO_SIN_MATCH_HOJA1": 141,
        "BLOQUEADO_SIN_REGISTROS_HASTA_2025": 5,
        "BLOQUEADO_MULTICARRERA_SIN_TRAZABILIDAD": 1,
    }
    add("PENDIENTES_POR_ESTADO_ESPERADOS", counts == expected, json.dumps(counts, ensure_ascii=False))
    add("ID_FILA_5809_UNICO", "ID_FILA_5809" in df.columns and df["ID_FILA_5809"].is_unique, str(df["ID_FILA_5809"].duplicated().sum()) if "ID_FILA_5809" in df.columns else "sin columna")
    add("CODIGO_UNICO_EXISTE", "CODIGO_UNICO" in df.columns, "CODIGO_UNICO")
    add("DOCUMENTO_EXISTE", {"NUM_DOCUMENTO", "DV"}.issubset(df.columns) or "RUT_NORM" in df.columns or "_RUT_NORM" in df.columns, "NUM_DOCUMENTO/DV o RUT_NORM")
    return validations


def classify_pending(row: pd.Series) -> dict[str, str]:
    estado = str(row.get("ESTADO_CALCULO", ""))
    if estado == "BLOQUEADO_SIN_MATCH_HOJA1":
        return {
            "TIPO_BLOQUEO_GESTION": "FALTA_HISTORIAL_ACADEMICO_HOJA1",
            "ACCION_REQUERIDA": "Buscar fuente academica equivalente por ramo/unidad.",
            "FUENTE_NECESARIA": "RUT, CODCARR/CODCLI, CODRAMO, ANO, PERIODO, DESCRIPCION_ESTADO.",
            "PUEDE_CALCULARSE_ACTUALMENTE": "NO",
            "MOTIVO_NO_CALCULO": "No existe match en Hoja1 con historial calculable.",
            "NIVEL_RESPALDO_ACTUAL": "PENDIENTE_SIN_HISTORIAL_HOJA1",
        }
    if estado == "BLOQUEADO_SIN_REGISTROS_HASTA_2025":
        return {
            "TIPO_BLOQUEO_GESTION": "FALTA_HISTORIAL_HASTA_2025",
            "ACCION_REQUERIDA": "Buscar historial academico 2025 o anterior.",
            "FUENTE_NECESARIA": "Historial academico con ANO <= 2025.",
            "PUEDE_CALCULARSE_ACTUALMENTE": "NO",
            "MOTIVO_NO_CALCULO": "Sin registros academicos hasta 2025 en fuente calculable previa.",
            "NIVEL_RESPALDO_ACTUAL": "PENDIENTE_SIN_HISTORIAL_HASTA_2025",
        }
    if estado == "BLOQUEADO_MULTICARRERA_SIN_TRAZABILIDAD":
        return {
            "TIPO_BLOQUEO_GESTION": "FALTA_DECISION_CODCARR",
            "ACCION_REQUERIDA": "Documentar equivalencia de carrera para ID 70 antes de calcular.",
            "FUENTE_NECESARIA": "Equivalencia institucional entre INGENIERIA EN LOGISTICA y CODCARR candidato.",
            "PUEDE_CALCULARSE_ACTUALMENTE": "NO",
            "MOTIVO_NO_CALCULO": "ID 70 con candidatos Hoja1 AUDT | ICDA sin equivalencia documentada.",
            "NIVEL_RESPALDO_ACTUAL": "PENDIENTE_MULTICARRERA_SIN_TRAZABILIDAD",
        }
    return {
        "TIPO_BLOQUEO_GESTION": "PENDIENTE_NO_CLASIFICADO",
        "ACCION_REQUERIDA": "Revisar estado.",
        "FUENTE_NECESARIA": "No determinada.",
        "PUEDE_CALCULARSE_ACTUALMENTE": "NO",
        "MOTIVO_NO_CALCULO": "Estado no reconocido.",
        "NIVEL_RESPALDO_ACTUAL": "REQUIERE_REVISION",
    }


def candidate_files() -> list[Path]:
    out = []
    for path in BASE.rglob("*"):
        if not path.is_file():
            continue
        if any(part in {".git", ".venv", "__pycache__"} for part in path.parts):
            continue
        if RUN_DIR in path.parents:
            continue
        if path.suffix.lower() not in SUFFIXES:
            continue
        rel_norm = norm_text(path.relative_to(BASE))
        if path == PROMEDIOS or any(norm_text(k) in rel_norm for k in KEYWORDS):
            out.append(path)
    return sorted(set(out))


def detect_columns(cols: list[str]) -> dict[str, str]:
    norm_map = {norm_text(c): c for c in cols}

    def first_exact(patterns: list[str]) -> str:
        for pat in patterns:
            original = norm_map.get(norm_text(pat))
            if original:
                return original
        return ""

    return {
        "rut": first_exact(["RUT_NORM", "RUT", "NUM_DOCUMENTO", "N_DOC", "DOCUMENTO"]),
        "dv": first_exact(["DV", "DIG", "DIGITO"]),
        "codigo_unico": first_exact(["CODIGO_UNICO", "CODIGO_UNICO_5809", "CODIGO_UNICO_SIES"]),
        "codcarr": first_exact(["CODCARR", "CODCARPR", "CODCARR_DECIDIDO", "CODCARR_CANDIDATO_HOJA1", "CODIGO_CARRERA"]),
        "codcli": first_exact(["CODCLI"]),
        "carrera": first_exact(["NOMBRE_CARRERA", "CARRERA", "NOMBRE_CARRERA_5810", "NOMBRE_CARRERA_MATRIZ"]),
        "codramo": first_exact(["CODRAMO", "RAMOEQUIV", "COD_ASIGNATURA", "CODIGO_ASIGNATURA", "ID_ASIGNATURA", "ASIGNATURA"]),
        "ano": first_exact(["ANO", "ANIO", "AÑO"]),
        "periodo": first_exact(["PERIODO", "SEMESTRE"]),
        "estado": first_exact(["DESCRIPCION_ESTADO", "ESTADO_ACADEMICO", "ESTADO"]),
        "nota": first_exact(["NOTA_FINAL", "NOTA", "EXAMEN"]),
        "convalidado": first_exact(["CONVALIDADO", "CONVALIDACION", "RECONOCIMIENTO", "HOMOLOG"]),
    }


def classify_source(mapping: dict[str, str]) -> str:
    has_id = bool(mapping["rut"])
    has_career = bool(mapping["codigo_unico"] or mapping["codcarr"] or mapping["codcli"] or mapping["carrera"])
    has_ramo = bool(mapping["codramo"])
    has_year = bool(mapping["ano"])
    has_period = bool(mapping["periodo"])
    has_state = bool(mapping["estado"])
    if has_id and has_career and has_ramo and has_year and has_period and has_state:
        return "FUENTE_ACADEMICA_CALCULABLE"
    if has_id or has_career or has_ramo or has_year or has_state:
        return "EVIDENCIA_AUXILIAR_NO_CALCULABLE"
    return "NO_RELEVANTE"


def inspect_sources() -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for path in candidate_files():
        suffix = path.suffix.lower()
        try:
            if suffix in {".xlsx", ".xlsm"}:
                xls = pd.ExcelFile(path)
                for sheet in xls.sheet_names:
                    try:
                        header = pd.read_excel(path, sheet_name=sheet, nrows=0, dtype=str)
                        mapping = detect_columns(list(header.columns))
                        cls = classify_source(mapping)
                        filas = ""
                        try:
                            wb = load_workbook(path, read_only=True, data_only=True)
                            ws = wb[sheet]
                            filas = max(ws.max_row - 1, 0)
                            wb.close()
                        except Exception:
                            pass
                        rows.append(source_row(path, suffix, sheet, filas, list(header.columns), mapping, cls, ""))
                    except Exception as exc:
                        rows.append(source_row(path, suffix, sheet, "", [], {}, "REQUIERE_REVISION", f"{type(exc).__name__}: {exc}"))
            elif suffix == ".parquet":
                df = pd.read_parquet(path)
                mapping = detect_columns(list(df.columns))
                rows.append(source_row(path, suffix, "", len(df), list(df.columns), mapping, classify_source(mapping), ""))
            else:
                df = read_csv_flexible(path, nrows=50)
                mapping = detect_columns(list(df.columns))
                filas = ""
                try:
                    with path.open("rb") as fh:
                        filas = max(sum(1 for _ in fh) - 1, 0)
                except Exception:
                    pass
                rows.append(source_row(path, suffix, "", filas, list(df.columns), mapping, classify_source(mapping), ""))
        except Exception as exc:
            rows.append(source_row(path, suffix, "", "", [], {}, "REQUIERE_REVISION", f"{type(exc).__name__}: {exc}"))
    return pd.DataFrame(rows)


def source_row(path: Path, suffix: str, sheet: str, rows: Any, cols: list[str], mapping: dict[str, str], cls: str, error: str) -> dict[str, Any]:
    return {
        "ruta": str(path),
        "tipo_archivo": suffix.lstrip("."),
        "hoja": sheet,
        "filas": rows,
        "columnas": len(cols),
        "columnas_relevantes_detectadas": " | ".join(v for v in mapping.values() if v),
        "tiene_RUT_documento": "SI" if mapping.get("rut") else "NO",
        "tiene_CODCARR_CODCLI_carrera": "SI" if (mapping.get("codigo_unico") or mapping.get("codcarr") or mapping.get("codcli") or mapping.get("carrera")) else "NO",
        "tiene_CODRAMO_ramo_asignatura": "SI" if mapping.get("codramo") else "NO",
        "tiene_ANO": "SI" if mapping.get("ano") else "NO",
        "tiene_PERIODO": "SI" if mapping.get("periodo") else "NO",
        "tiene_DESCRIPCION_ESTADO_estado": "SI" if mapping.get("estado") else "NO",
        "tiene_nota": "SI" if mapping.get("nota") else "NO",
        "tiene_convalidado": "SI" if mapping.get("convalidado") else "NO",
        "clasificacion": cls,
        "col_rut": mapping.get("rut", ""),
        "col_dv": mapping.get("dv", ""),
        "col_codigo_unico": mapping.get("codigo_unico", ""),
        "col_codcarr": mapping.get("codcarr", ""),
        "col_codcli": mapping.get("codcli", ""),
        "col_carrera": mapping.get("carrera", ""),
        "col_codramo": mapping.get("codramo", ""),
        "col_ano": mapping.get("ano", ""),
        "col_periodo": mapping.get("periodo", ""),
        "col_estado": mapping.get("estado", ""),
        "error": error,
    }


def load_candidate_source(row: pd.Series) -> pd.DataFrame:
    path = Path(row["ruta"])
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return pd.read_excel(path, sheet_name=row["hoja"], dtype=str, keep_default_na=False)
    if suffix == ".parquet":
        return pd.read_parquet(path).astype(str)
    return read_csv_flexible(path)


def normalize_candidate(df: pd.DataFrame, source: pd.Series) -> pd.DataFrame:
    out = df.copy()
    rut_col = source.get("col_rut", "")
    dv_col = source.get("col_dv", "")
    if rut_col:
        if dv_col and dv_col in out.columns and rut_col in out.columns and rut_col != dv_col:
            out["_RUT_NORM_SRC"] = [norm_rut_body_dv(r, d) for r, d in zip(out[rut_col], out[dv_col])]
        else:
            out["_RUT_NORM_SRC"] = out[rut_col].map(split_rut)
    else:
        out["_RUT_NORM_SRC"] = ""
    for key in ["col_codigo_unico", "col_codcarr", "col_codcli", "col_carrera", "col_codramo", "col_estado"]:
        col = source.get(key, "")
        out[f"_{key.upper()}_SRC"] = out[col].astype(str) if col and col in out.columns else ""
    ano_col = source.get("col_ano", "")
    periodo_col = source.get("col_periodo", "")
    out["_ANO_NUM"] = pd.to_numeric(out[ano_col], errors="coerce") if ano_col and ano_col in out.columns else pd.NA
    out["_PERIODO_NUM"] = pd.to_numeric(out[periodo_col], errors="coerce") if periodo_col and periodo_col in out.columns else pd.NA
    out["_CODIGO_UNICO_NORM_SRC"] = out["_COL_CODIGO_UNICO_SRC"].map(norm_text)
    out["_CODCARR_NORM_SRC"] = out["_COL_CODCARR_SRC"].map(norm_text)
    out["_CARRERA_NORM_SRC"] = out["_COL_CARRERA_SRC"].map(norm_text)
    return out


def expected_5810(df5810: pd.DataFrame) -> dict[str, dict[str, str]]:
    return {
        row["CODIGO_UNICO"]: {
            "NOMBRE_CARRERA_5810": row.get("NOMBRE_CARRERA", ""),
            "CODIGO_UNICO": row.get("CODIGO_UNICO", ""),
            "PLAN_ESTUDIOS": row.get("PLAN_ESTUDIOS", ""),
        }
        for _, row in df5810.iterrows()
    }


def resolve_career(hist: pd.DataFrame, pending: pd.Series, exp: dict[str, str]) -> dict[str, str]:
    codigo = str(pending["CODIGO_UNICO"])
    codigo_norm = norm_text(codigo)
    expected_name = exp.get(codigo, {}).get("NOMBRE_CARRERA_5810", "")
    expected_norm = norm_text(expected_name)

    if "_CODIGO_UNICO_NORM_SRC" in hist.columns and hist["_CODIGO_UNICO_NORM_SRC"].astype(str).str.len().gt(0).any():
        exact = hist[hist["_CODIGO_UNICO_NORM_SRC"].eq(codigo_norm)]
        if not exact.empty:
            return {
                "estado": "RESUELTO_CARRERA_UNICA",
                "hist_indices": "|".join(map(str, exact.index.tolist())),
                "codcarr_decidido": "|".join(sorted(set(exact["_COL_CODCARR_SRC"].astype(str).str.strip()) - {""})),
                "motivo": "CODIGO_UNICO exacto en fuente alternativa.",
                "nivel": "FUENTE_ACADEMICA_EQUIVALENTE_CODIGO_UNICO",
            }
        return {"estado": "NO_CALCULABLE_SIN_CARRERA_COMPATIBLE", "hist_indices": "", "codcarr_decidido": "", "motivo": "Fuente contiene CODIGO_UNICO, pero no coincide con 5809.", "nivel": "CONTRASTE_CODIGO_UNICO"}

    if expected_norm and "_CARRERA_NORM_SRC" in hist.columns and hist["_CARRERA_NORM_SRC"].astype(str).str.len().gt(0).any():
        exact_name = hist[hist["_CARRERA_NORM_SRC"].eq(expected_norm)]
        if not exact_name.empty:
            cods = sorted(set(exact_name["_COL_CODCARR_SRC"].astype(str).str.strip()) - {""})
            return {
                "estado": "RESUELTO_DECISION_TECNICA_DOCUMENTADA" if len(cods) > 1 else "RESUELTO_CARRERA_UNICA",
                "hist_indices": "|".join(map(str, exact_name.index.tolist())),
                "codcarr_decidido": "|".join(cods),
                "motivo": "Nombre de carrera de fuente coincide exactamente con 5810.",
                "nivel": "DECISION_INTERNA_TECNICA_NOMBRE_5810",
            }

    cods = sorted(set(hist["_COL_CODCARR_SRC"].astype(str).str.strip()) - {""}) if "_COL_CODCARR_SRC" in hist.columns else []
    if len(cods) == 1 and expected_norm and hist["_CARRERA_NORM_SRC"].eq(expected_norm).any():
        return {
            "estado": "RESUELTO_CARRERA_UNICA",
            "hist_indices": "|".join(map(str, hist.index.tolist())),
            "codcarr_decidido": cods[0],
            "motivo": "Unico CODCARR con nombre de carrera compatible.",
            "nivel": "FUENTE_ACADEMICA_EQUIVALENTE_CARRERA_UNICA",
        }
    if len(cods) > 1:
        return {"estado": "PENDIENTE_MULTICARRERA_SIN_TRAZABILIDAD", "hist_indices": "", "codcarr_decidido": "|".join(cods), "motivo": "Multiples CODCARR sin decision inequívoca.", "nivel": "PENDIENTE"}
    return {"estado": "NO_CALCULABLE_SIN_CARRERA_COMPATIBLE", "hist_indices": "", "codcarr_decidido": "|".join(cods), "motivo": "No hay evidencia de compatibilidad con CODIGO_UNICO/5810.", "nivel": "PENDIENTE"}


def build_hist_for_function(hist: pd.DataFrame, pending: pd.Series, source: pd.Series) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "RUT_NORM": pending["_RUT_NORM"],
            "CODIGO_UNICO": pending["CODIGO_UNICO"],
            "ANO": pd.to_numeric(hist["_ANO_NUM"], errors="coerce").astype("Int64"),
            "PERIODO": pd.to_numeric(hist["_PERIODO_NUM"], errors="coerce").astype("Int64"),
            "CODRAMO": hist["_COL_CODRAMO_SRC"].astype(str),
            "DESCRIPCION_ESTADO": hist["_COL_ESTADO_SRC"].astype(str),
        }
    ).dropna(subset=["ANO", "PERIODO"])


def main() -> int:
    for directory in [RESULTS, AUDIT, REPORTS]:
        directory.mkdir(parents=True, exist_ok=True)

    missing = [p for p in required_files() if not p.exists()]
    if missing:
        msg = "\n".join([f"BLOQUEO: falta archivo obligatorio {p}" for p in missing])
        (REPORTS / "BLOQUEO_FALTA_ARCHIVO.md").write_text(msg + "\n", encoding="utf-8")
        print(msg)
        return 2

    hashes = pd.DataFrame(
        [
            {
                "tipo": label,
                "ruta": str(path),
                "sha256": sha256(path),
                "tamano_bytes": path.stat().st_size,
                "fecha_modificacion": datetime.fromtimestamp(path.stat().st_mtime, TZ).isoformat(timespec="seconds"),
            }
            for label, path in [
                ("codigo_gobernanza_v2", CODIGO_GOB),
                ("matriz_vigente_2224_147", MATRIX_IN),
                ("precarga_5809", PRECARGA_5809),
                ("precarga_5810", PRECARGA_5810),
                ("promedios_7804", PROMEDIOS),
                ("instructivo", INSTRUCTIVO),
            ]
        ]
    )
    write_csv(AUDIT / "hashes_entradas.csv", hashes)

    write_json(
        RUN_DIR / "manifest_continuacion_147_pendientes.json",
        {
            "fecha_hora": now_iso(),
            "proceso": "Avance Curricular SIES 2026",
            "subproyecto": "Matricula 5809 / columnas 16 a 21",
            "anio_referencia": 2025,
            "rutas_entrada": hashes.to_dict(orient="records"),
            "usuario": os.environ.get("USER", ""),
            "cwd": str(BASE),
            "csv_carga_generado": False,
            "sies_ready_generado": False,
            "fuentes_originales_modificadas": False,
        },
    )

    construir_resumen_historico = load_function()
    matrix = load_matrix()
    initial_validations = validate_initial_matrix(matrix)
    if any(v["RESULTADO"] != "OK" for v in initial_validations):
        val_df = pd.DataFrame(initial_validations)
        write_csv(AUDIT / "validaciones_finales.csv", val_df)
        write_xlsx(RESULTS / "BLOQUEO_VALIDACION_INICIAL.xlsx", {"VALIDACIONES": val_df})
        print("BLOQUEO DETECTADO")
        print("Motivo: validacion inicial de matriz 2224/147 fallo.")
        print(f"Archivo: {MATRIX_IN}")
        print("Hoja: MATRIZ_TRABAJO_2371")
        print("Columna: ver validaciones_finales.csv")
        print("Acción recomendada: revisar cifras de entrada antes de continuar.")
        print("No se generó matriz final de trabajo.")
        return 3

    matrix["_RUT_NORM"] = [norm_rut_body_dv(r, d) for r, d in zip(matrix["NUM_DOCUMENTO"], matrix["DV"])]
    pending_mask = matrix["ESTADO_PREPARACION_16_21"].eq("PENDIENTE_NO_CARGA")
    pendientes = matrix[pending_mask].copy()
    extra = pd.DataFrame([classify_pending(r) for _, r in pendientes.iterrows()])
    pendientes_base = pd.concat([pendientes.reset_index(drop=True), extra], axis=1)

    inventory = inspect_sources()
    write_csv(AUDIT / "inventario_fuentes_candidatas.csv", inventory)

    df5810 = load_5810()
    exp_map = expected_5810(df5810)
    calculable = inventory[inventory["clasificacion"].eq("FUENTE_ACADEMICA_CALCULABLE")].copy()
    diagnostics: list[dict[str, Any]] = []
    decisions: list[dict[str, Any]] = []
    new_rows: list[dict[str, Any]] = []
    detail_rows: list[dict[str, Any]] = []

    updated = matrix.copy()
    already_new: set[str] = set()

    for _, source in calculable.iterrows():
        try:
            src = load_candidate_source(source)
            srcn = normalize_candidate(src, source)
        except Exception as exc:
            diagnostics.append(
                {
                    "ID_FILA_5809": "",
                    "fuente_usada": source["ruta"],
                    "hoja_usada": source["hoja"],
                    "decision_calculo": "NO_CALCULABLE_FALTAN_CAMPOS",
                    "detalle": f"No se pudo leer fuente: {type(exc).__name__}: {exc}",
                }
            )
            continue
        srcn = srcn[srcn["_RUT_NORM_SRC"].astype(str).str.len().gt(0)].copy()
        srcn["_ANO_NUM"] = pd.to_numeric(srcn["_ANO_NUM"], errors="coerce")
        srcn["_PERIODO_NUM"] = pd.to_numeric(srcn["_PERIODO_NUM"], errors="coerce")
        src_le = srcn[srcn["_ANO_NUM"].le(2025)].copy()

        for _, p in pendientes.iterrows():
            pid = str(p["ID_FILA_5809"])
            rut = str(p["_RUT_NORM"])
            stud_all = srcn[srcn["_RUT_NORM_SRC"].eq(rut)]
            stud = src_le[src_le["_RUT_NORM_SRC"].eq(rut)]
            decision = "NO_CALCULABLE_SIN_RUT"
            career_resolution = {"estado": "", "hist_indices": "", "codcarr_decidido": "", "motivo": "", "nivel": ""}
            hist_use = pd.DataFrame()
            if not stud_all.empty and stud.empty:
                decision = "NO_CALCULABLE_SIN_HISTORIAL_2025"
            elif not stud.empty:
                career_resolution = resolve_career(stud, p, exp_map)
                if pid == "70":
                    cods_id70 = [c for c in career_resolution.get("codcarr_decidido", "").split("|") if c]
                    if len(cods_id70) != 1:
                        career_resolution = {
                            "estado": "PENDIENTE_MULTICARRERA_SIN_TRAZABILIDAD",
                            "hist_indices": "",
                            "codcarr_decidido": "|".join(cods_id70),
                            "motivo": "ID 70 no se resuelve automaticamente: requiere equivalencia documentada para AUDT | ICDA.",
                            "nivel": "PENDIENTE_REGLA_EXPLICITA_ID70",
                        }
                if career_resolution["estado"] in {"RESUELTO_CARRERA_UNICA", "RESUELTO_DECISION_TECNICA_DOCUMENTADA"}:
                    idx = [int(x) for x in career_resolution["hist_indices"].split("|") if x]
                    hist_use = stud.loc[idx].copy() if idx else pd.DataFrame()
                    if hist_use.empty:
                        decision = "REQUIERE_REVISION"
                    else:
                        decision = "CALCULABLE"
                elif career_resolution["estado"] == "PENDIENTE_MULTICARRERA_SIN_TRAZABILIDAD":
                    decision = "NO_CALCULABLE_MULTICARRERA_SIN_DECISION"
                else:
                    decision = "NO_CALCULABLE_SIN_CARRERA_COMPATIBLE"

            years = sorted(set(pd.to_numeric(stud["_ANO_NUM"], errors="coerce").dropna().astype(int).astype(str))) if not stud.empty else []
            periods = sorted(set(pd.to_numeric(stud["_PERIODO_NUM"], errors="coerce").dropna().astype(int).astype(str))) if not stud.empty else []
            states = sorted(set(stud["_COL_ESTADO_SRC"].astype(str).str.strip()) - {""})[:20] if not stud.empty else []
            cods = sorted(set(stud["_COL_CODCARR_SRC"].astype(str).str.strip()) - {""}) if not stud.empty and "_COL_CODCARR_SRC" in stud.columns else []

            diagnostics.append(
                {
                    "ID_FILA_5809": pid,
                    "RUT_NORM": rut,
                    "CODIGO_UNICO": p["CODIGO_UNICO"],
                    "encontro_estudiante": "SI" if not stud_all.empty else "NO",
                    "encontro_historial_hasta_2025": "SI" if not stud.empty else "NO",
                    "encontro_carrera_compatible": "SI" if decision == "CALCULABLE" else "NO",
                    "cantidad_ramos_unidades": hist_use["_COL_CODRAMO_SRC"].nunique() if not hist_use.empty else 0,
                    "anios_presentes": " | ".join(years),
                    "periodos_presentes": " | ".join(periods),
                    "estados_presentes": " | ".join(states),
                    "codcarr_candidatos": " | ".join(cods),
                    "fuente_usada": source["ruta"],
                    "hoja_usada": source["hoja"],
                    "decision_calculo": decision,
                    "detalle": career_resolution.get("motivo", ""),
                }
            )

            if decision != "CALCULABLE" or pid in already_new:
                continue

            func_hist = build_hist_for_function(hist_use, p, source)
            func_hist = func_hist[pd.to_numeric(func_hist["ANO"], errors="coerce").le(2025)].copy()
            if func_hist.empty:
                continue
            resumen = construir_resumen_historico(func_hist)
            if resumen.empty:
                continue
            res = resumen.iloc[0].to_dict()
            row_idx = updated.index[updated["ID_FILA_5809"].astype(str).eq(pid)][0]
            for field in TARGET_FIELDS:
                updated.at[row_idx, f"{field}_PROPUESTO"] = str(res[field])
                calc_col = f"{field}_CALC"
                if calc_col in updated.columns:
                    updated.at[row_idx, calc_col] = str(res[field])
            estado_calc = "CALCULADO_DECISION_TECNICA_DOCUMENTADA" if career_resolution["estado"] == "RESUELTO_DECISION_TECNICA_DOCUMENTADA" else "CALCULADO_FUENTE_ALTERNATIVA_VALIDADA"
            updated.at[row_idx, "ESTADO_CALCULO"] = estado_calc
            updated.at[row_idx, "ESTADO_PREPARACION_16_21"] = "CALCULADO_NO_CARGA"
            updated.at[row_idx, "METODO_FILTRO_CARRERA"] = career_resolution["estado"]
            updated.at[row_idx, "CODCARR_USADO"] = career_resolution["codcarr_decidido"]
            updated.at[row_idx, "CODCLI_USADOS"] = " | ".join(sorted(set(hist_use["_COL_CODCLI_SRC"].astype(str).str.strip()) - {""}))
            updated.at[row_idx, "ANIOS_USADOS"] = " | ".join(sorted(set(func_hist["ANO"].astype(str))))
            updated.at[row_idx, "PERIODOS_USADOS"] = " | ".join(sorted(set(func_hist["PERIODO"].astype(str))))
            updated.at[row_idx, "REGISTROS_HOJA1_FILTRADOS"] = str(len(func_hist))
            updated.at[row_idx, "FUENTE_CALCULO_16_21"] = source["ruta"]
            updated.at[row_idx, "HOJA_CALCULO_16_21"] = source["hoja"]
            updated.at[row_idx, "NIVEL_RESPALDO_CALCULO"] = career_resolution["nivel"]
            already_new.add(pid)
            new_row = updated.loc[row_idx].to_dict()
            new_rows.append(new_row)
            decisions.append(
                {
                    "ID_FILA_5809": pid,
                    "RUT_NORM": rut,
                    "CODIGO_UNICO": p["CODIGO_UNICO"],
                    "SIES_CARRERA": re.search(r"C([0-9]+)", str(p["CODIGO_UNICO"])).group(1) if re.search(r"C([0-9]+)", str(p["CODIGO_UNICO"])) else "",
                    "NOMBRE_CARRERA_5810": exp_map.get(str(p["CODIGO_UNICO"]), {}).get("NOMBRE_CARRERA_5810", ""),
                    "NOMBRE_CARRERA_MATRIZ": "",
                    "CODCARR_CANDIDATOS": " | ".join(cods),
                    "CODCARR_DECIDIDO": career_resolution["codcarr_decidido"],
                    "MOTIVO_DECISION": career_resolution["motivo"],
                    "NIVEL_RESPALDO": career_resolution["nivel"],
                    "FUENTE_DECISION": source["ruta"] + ("::" + str(source["hoja"]) if source["hoja"] else ""),
                    "OBSERVACION": "DECISION_INTERNA_TECNICA" if "DECISION_INTERNA" in career_resolution["nivel"] else "",
                }
            )
            for _, h in hist_use.iterrows():
                detail = {
                    "ID_FILA_5809": pid,
                    "RUT_NORM": rut,
                    "CODIGO_UNICO": p["CODIGO_UNICO"],
                    "FUENTE": source["ruta"],
                    "HOJA": source["hoja"],
                    "CODCARR_USADO": career_resolution["codcarr_decidido"],
                    "ANO": h["_ANO_NUM"],
                    "PERIODO": h["_PERIODO_NUM"],
                    "CODRAMO": h["_COL_CODRAMO_SRC"],
                    "DESCRIPCION_ESTADO": h["_COL_ESTADO_SRC"],
                    "CODCLI": h["_COL_CODCLI_SRC"],
                }
                detail_rows.append(detail)

    diag_df = pd.DataFrame(diagnostics)
    decisions_df = pd.DataFrame(decisions)
    new_df = pd.DataFrame(new_rows)
    detail_df = pd.DataFrame(detail_rows)

    write_csv(AUDIT / "diagnostico_cruce_147_fuentes.csv", diag_df)
    if not detail_df.empty:
        write_csv(AUDIT / "detalle_historial_usado_nuevos.csv", detail_df)
    else:
        write_csv(AUDIT / "detalle_historial_usado_nuevos.csv", pd.DataFrame(columns=["ID_FILA_5809", "RUT_NORM", "CODIGO_UNICO", "FUENTE", "HOJA", "CODCARR_USADO", "ANO", "PERIODO", "CODRAMO", "DESCRIPCION_ESTADO", "CODCLI"]))

    total_calc = int(updated["ESTADO_PREPARACION_16_21"].eq("CALCULADO_NO_CARGA").sum())
    total_pend = int(updated["ESTADO_PREPARACION_16_21"].eq("PENDIENTE_NO_CARGA").sum())
    nuevos = len(already_new)

    final_validations: list[dict[str, str]] = initial_validations.copy()

    def add_val(name: str, ok: bool, detail: str) -> None:
        final_validations.append({"VALIDACION": name, "RESULTADO": "OK" if ok else "ERROR", "DETALLE": detail})

    add_val("TOTAL_FILAS_FINAL_2371", len(updated) == 2371, str(len(updated)))
    add_val("ID_FILA_5809_UNICO_FINAL", updated["ID_FILA_5809"].is_unique, str(updated["ID_FILA_5809"].duplicated().sum()))
    add_val("NO_PERDIDA_COLUMNAS", set(matrix.columns).issubset(set(updated.columns)), f"inicial={len(matrix.columns)} final={len(updated.columns)}")
    add_val("CALCULADOS_MAS_PENDIENTES_2371", total_calc + total_pend == 2371, f"{total_calc}+{total_pend}")
    add_val("CALCULADOS_ANTERIORES_SIGUEN", total_calc >= 2224, str(total_calc))
    old_pending_ids = set(pendientes["ID_FILA_5809"].astype(str))
    add_val("NUEVOS_SOLO_DESDE_PENDIENTES", set(already_new).issubset(old_pending_ids), " | ".join(sorted(already_new)))
    add_val("NINGUN_NUEVO_SIN_DETALLE", nuevos == 0 or set(detail_df["ID_FILA_5809"].astype(str)) == set(already_new), f"nuevos={nuevos} detalle_ids={detail_df['ID_FILA_5809'].nunique() if not detail_df.empty else 0}")
    if nuevos:
        calc_new = updated[updated["ID_FILA_5809"].astype(str).isin(already_new)]
        add_val("NUEVOS_CODCARR_USADO", calc_new["CODCARR_USADO"].astype(str).str.strip().ne("").all(), str(nuevos))
        add_val("NUEVOS_FUENTE_CALCULO", calc_new.get("FUENTE_CALCULO_16_21", pd.Series([""] * len(calc_new))).astype(str).str.strip().ne("").all(), str(nuevos))
        add_val("NUEVOS_REGISTROS_USADOS", pd.to_numeric(calc_new["REGISTROS_HOJA1_FILTRADOS"], errors="coerce").gt(0).all(), str(nuevos))
        add_val("NUEVOS_SI_NO", calc_new[["CURSO_1ER_SEM_PROPUESTO", "CURSO_2DO_SEM_PROPUESTO"]].isin(["SI", "NO"]).all().all(), str(nuevos))
        for col in ["UNIDADES_CURSADAS_PROPUESTO", "UNIDADES_APROBADAS_PROPUESTO", "UNID_CURSADAS_TOTAL_PROPUESTO", "UNID_APROBADAS_TOTAL_PROPUESTO"]:
            nums = pd.to_numeric(calc_new[col], errors="coerce")
            add_val(f"NUEVOS_{col}_ENTERO_NO_NEGATIVO", nums.notna().all() and nums.ge(0).all() and (nums % 1).eq(0).all(), col)
        uc = pd.to_numeric(calc_new["UNIDADES_CURSADAS_PROPUESTO"], errors="coerce")
        ua = pd.to_numeric(calc_new["UNIDADES_APROBADAS_PROPUESTO"], errors="coerce")
        uct = pd.to_numeric(calc_new["UNID_CURSADAS_TOTAL_PROPUESTO"], errors="coerce")
        uat = pd.to_numeric(calc_new["UNID_APROBADAS_TOTAL_PROPUESTO"], errors="coerce")
        add_val("NUEVOS_APROBADAS_MENOR_IGUAL_CURSADAS", ua.le(uc).all(), str(nuevos))
        add_val("NUEVOS_APROBADAS_TOTAL_MENOR_IGUAL_CURSADAS_TOTAL", uat.le(uct).all(), str(nuevos))
        max_year = pd.to_numeric(detail_df["ANO"], errors="coerce").max()
        add_val("NUEVOS_ANIOS_USADOS_HASTA_2025", pd.isna(max_year) or max_year <= 2025, str(max_year))
    add_val("CSV_CARGA_GENERADO_NO", True, "NO")
    add_val("SIES_READY_GENERADO_NO", True, "NO")
    add_val("FUENTES_ORIGINALES_MODIFICADAS_NO", True, "NO")

    val_df = pd.DataFrame(final_validations)
    write_csv(AUDIT / "validaciones_finales.csv", val_df)
    if (val_df["RESULTADO"] == "ERROR").any():
        write_xlsx(RESULTS / "BLOQUEO_VALIDACIONES_FINALES.xlsx", {"VALIDACIONES": val_df, "DIAGNOSTICO": diag_df})
        print("BLOQUEO DETECTADO")
        print("Motivo: validaciones finales con ERROR.")
        print(f"Archivo: {MATRIX_IN}")
        print("Hoja: MATRIZ_TRABAJO_2371")
        print("Columna: ver validaciones_finales.csv")
        print("Acción recomendada: revisar errores antes de generar matriz final.")
        print("No se generó matriz final de trabajo.")
        return 4

    updated_for_output = updated.drop(columns=["_RUT_NORM"], errors="ignore")
    pendientes_restantes = updated_for_output[updated_for_output["ESTADO_PREPARACION_16_21"].eq("PENDIENTE_NO_CARGA")].copy()
    calculados_total = updated_for_output[updated_for_output["ESTADO_PREPARACION_16_21"].eq("CALCULADO_NO_CARGA")].copy()
    pendientes_base_out = pendientes_base.drop(columns=["_RUT_NORM"], errors="ignore")

    base_name = f"MATRIZ_TRABAJO_ACTUALIZADA_{total_calc}_CALCULADOS_{total_pend}_PENDIENTES_NO_CARGA"
    xlsx_out = RESULTS / f"{base_name}.xlsx"
    tsv_out = RESULTS / f"{base_name}.tsv"
    write_tsv(tsv_out, updated_for_output)

    resumen = pd.DataFrame(
        [
            {"METRICA": "TOTAL_5809", "VALOR": len(updated_for_output)},
            {"METRICA": "CALCULADOS_ANTERIORES", "VALOR": 2224},
            {"METRICA": "NUEVOS_CALCULADOS", "VALOR": nuevos},
            {"METRICA": "CALCULADOS_TOTAL", "VALOR": total_calc},
            {"METRICA": "PENDIENTES_RESTANTES", "VALOR": total_pend},
            {"METRICA": "CSV_CARGA_GENERADO", "VALOR": "NO"},
            {"METRICA": "SIES_READY_GENERADO", "VALOR": "NO"},
            {"METRICA": "FUENTES_ORIGINALES_MODIFICADAS", "VALOR": "NO"},
        ]
    )
    estados = updated_for_output["ESTADO_CALCULO"].value_counts().rename_axis("ESTADO_CALCULO").reset_index(name="N")
    pend_resumen = pendientes_restantes["ESTADO_CALCULO"].value_counts().rename_axis("ESTADO_PENDIENTE").reset_index(name="N")
    criterio = pd.DataFrame(
        [
            {"TIPO": "REGLA_OFICIAL", "DESCRIPCION": "Columnas 16-21 se calculan desde actividad academica 2025 e historico hasta 2025.", "APLICACION": "No se calcula sin historial por ramo/unidad."},
            {"TIPO": "IMPLEMENTACION_TECNICA", "DESCRIPCION": "Se usa construir_resumen_historico(hist_mapeado) tras filtrar ANO <= 2025.", "APLICACION": "Solo para fuentes academicas calculables y carrera resuelta."},
            {"TIPO": "DECISION_INTERNA_TECNICA", "DESCRIPCION": "CODIGO_UNICO SIES es rector; no se decide ID 70 sin equivalencia documentada.", "APLICACION": "Pendiente si carrera no se aisla."},
        ]
    )
    fuentes_sheet = hashes.copy()
    sheets = {
        "RESUMEN": resumen,
        "ESTADOS_CALCULO": estados,
        "PENDIENTES_RESUMEN": pend_resumen,
        "VALIDACIONES": val_df,
        "CRITERIO": criterio,
        "INVENTARIO_FUENTES": inventory,
        "DIAGNOSTICO_CRUCE_147": diag_df,
        "DECISIONES_CARRERA": decisions_df,
        "MATRIZ_TRABAJO_2371": updated_for_output,
        "CALCULADOS_TOTAL": calculados_total,
        "PENDIENTES_RESTANTES": pendientes_restantes,
        "NUEVOS_CALCULADOS": new_df.drop(columns=["_RUT_NORM"], errors="ignore") if not new_df.empty else pd.DataFrame(columns=updated_for_output.columns),
        "DETALLE_HISTORIAL_USADO_NUEVOS": detail_df,
        "PENDIENTES_SIN_MATCH_HOJA1": pendientes_restantes[pendientes_restantes["ESTADO_CALCULO"].eq("BLOQUEADO_SIN_MATCH_HOJA1")],
        "PENDIENTES_SIN_HIST_2025": pendientes_restantes[pendientes_restantes["ESTADO_CALCULO"].eq("BLOQUEADO_SIN_REGISTROS_HASTA_2025")],
        "PENDIENTE_ID70": pendientes_restantes[pendientes_restantes["ID_FILA_5809"].astype(str).eq("70")],
        "PENDIENTES_147_BASE": pendientes_base_out,
        "FUENTES": fuentes_sheet,
    }
    write_xlsx(xlsx_out, sheets)

    report_lines = [
        "# Continuacion 147 pendientes columnas 16-21",
        "",
        f"Archivo de entrada: `{MATRIX_IN}`",
        f"Carpeta de salida: `{RUN_DIR}`",
        "",
        f"Total 5809: {len(updated_for_output)}",
        "Calculados anteriores: 2224",
        f"Nuevos calculados: {nuevos}",
        f"Calculados totales: {total_calc}",
        f"Pendientes restantes: {total_pend}",
        "",
        "Pendientes por tipo:",
    ]
    for _, r in pend_resumen.iterrows():
        report_lines.append(f"- {r['ESTADO_PENDIENTE']}: {r['N']}")
    report_lines += [
        "",
        f"Fuentes academicas calculables inventariadas: {len(calculable)}",
        "CSV de carga generado: NO",
        "SIES_READY generado: NO",
        "Fuentes originales modificadas: NO",
    ]
    (REPORTS / "INFORME_CONTINUACION_147_PENDIENTES_COLUMNAS_16_21_NO_CARGA.md").write_text("\n".join(report_lines) + "\n", encoding="utf-8")
    (REPORTS / "RESUMEN_EJECUCION.txt").write_text("\n".join(report_lines) + "\n", encoding="utf-8")

    shutil.copy2(Path(__file__).resolve(), AUDIT / Path(__file__).name)

    outputs = []
    for path in sorted(RUN_DIR.rglob("*")):
        if path.is_file():
            outputs.append({"ruta": str(path), "sha256": sha256(path), "tamano_bytes": path.stat().st_size})
    manifest_path = RUN_DIR / "manifest_continuacion_147_pendientes.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest.update(
        {
            "archivo_salida_principal": str(xlsx_out),
            "total_5809": len(updated_for_output),
            "calculados_anteriores": 2224,
            "nuevos_calculados": nuevos,
            "calculados_totales": total_calc,
            "pendientes_restantes": total_pend,
            "pendientes_por_tipo": pend_resumen.to_dict(orient="records"),
            "validaciones_ok": int((val_df["RESULTADO"] == "OK").sum()),
            "validaciones_revisar_error": int((val_df["RESULTADO"] != "OK").sum()),
            "archivos_generados": outputs,
            "csv_carga_generado": False,
            "sies_ready_generado": False,
            "fuentes_originales_modificadas": False,
        }
    )
    write_json(manifest_path, manifest)

    print("Proceso: Avance Curricular SIES 2026")
    print("Subproyecto: Matricula 5809 / columnas 16 a 21")
    print("Año referencia: 2025")
    print(f"Archivo de entrada principal: {MATRIX_IN}")
    print(f"Carpeta de salida: {RUN_DIR}")
    print(f"Total 5809: {len(updated_for_output)}")
    print("Calculados anteriores: 2224")
    print(f"Nuevos calculados: {nuevos}")
    print(f"Calculados totales: {total_calc}")
    print(f"Pendientes restantes: {total_pend}")
    print("Pendientes por tipo:")
    for _, r in pend_resumen.iterrows():
        print(f"- {r['ESTADO_PENDIENTE']}: {r['N']}")
    print(f"Validaciones OK: {(val_df['RESULTADO'] == 'OK').sum()} / REVISAR-ERROR: {(val_df['RESULTADO'] != 'OK').sum()}")
    print(f"Archivos generados: {len(outputs)}")
    print("CSV de carga generado: NO")
    print("SIES_READY generado: NO")
    print("Fuentes originales modificadas: NO")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
