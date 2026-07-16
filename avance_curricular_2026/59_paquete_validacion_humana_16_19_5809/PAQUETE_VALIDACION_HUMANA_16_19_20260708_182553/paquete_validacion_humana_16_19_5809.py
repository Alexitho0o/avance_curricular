#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Paquete de validacion humana 16-19 para Avance Curricular 5809.

No modifica fuentes originales, no corrige datos, no genera SIES_READY, no
recalcula 20-21 y no crea archivo de carga. Solo prepara expediente para
revision humana de los 423 casos pendientes.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Paquete validacion humana 16-19 5809"
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
NO_CORRECCION = "NO_APLICAR_CORRECCION_AUTOMATICA"
ESTADO_VALIDADOR_REQUERIDO = "APTO_PARA_CONTROL_PREVIO"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_59 = REPO / "avance_curricular_2026" / "59_paquete_validacion_humana_16_19_5809"

GOBERNANZA_51 = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)
VALIDADOR_52_SCRIPT = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "validador_gobernanza_precalculo_5809.py"
)
VALIDADOR_52_XLSX = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx"
)
CONSOLIDADO_58 = (
    REPO
    / "avance_curricular_2026/58_consolidado_resolucion_16_19_gobernado_5809/"
    / "CONSOLIDADO_RESOLUCION_16_19_GOBERNADO_20260708_181000/"
    / "CONSOLIDADO_RESOLUCION_16_19_GOBERNADO_5809.xlsx"
)
CONSOLIDADO_58_MANIFEST = CONSOLIDADO_58.with_name("manifest_consolidado_resolucion_16_19_gobernado_5809.json")
MATRIZ_53 = (
    REPO
    / "avance_curricular_2026/53_matriz_resolucion_bloqueos_16_19_gobernada_5809/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_20260708_170521/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_5809.xlsx"
)
PRECARGA_5809 = (
    REPO
    / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/"
    / "5809_Precarga Matrícula Avance Curricular 2026.csv"
)

OUTPUT_FILES = {
    "excel": "PAQUETE_VALIDACION_HUMANA_16_19_5809.xlsx",
    "informe": "INFORME_PAQUETE_VALIDACION_HUMANA_16_19_5809.md",
    "manifest": "manifest_paquete_validacion_humana_16_19_5809.json",
    "script": "paquete_validacion_humana_16_19_5809.py",
    "correo": "BORRADOR_CORREO_VALIDACION_16_19_5809.md",
}

SHEETS = OrderedDict(
    [
        ("00_DICTAMEN", "00_DICTAMEN"),
        ("01_VALIDACION_GOBERNANZA", "01_VALIDACION_GOBERNANZA"),
        ("02_RESUMEN_EJECUTIVO", "02_RESUMEN_EJECUTIVO"),
        ("03_DECISION_FUNCIONAL_256", "03_DECISION_FUNCIONAL_256"),
        ("04_MAPEO_INSTITUCIONAL_102", "04_MAPEO_INSTITUCIONAL_102"),
        ("05_FUENTE_COMPLEMENTARIA_57", "05_FUENTE_COMPLEMENTARIA_57"),
        ("06_REVISION_FUNCIONAL_8", "06_REVISION_FUNCIONAL_8"),
        ("07_PREGUNTAS_PARA_VALIDAR", "07_PREGUNTAS_PARA_VALIDAR"),
        ("08_CAMPOS_REQUERIDOS_RESPUESTA", "08_CAMPOS_REQUERIDOS_RESPUESTA"),
        ("09_PRIORIDAD_REVISION", "09_PRIORIDAD_REVISION"),
        ("10_NO_CORREGIR_AUTOMATICO", "10_NO_CORREGIR_AUTOMATICO"),
        ("11_FUENTES", "11_FUENTES"),
    ]
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def display(value: Any, max_len: int = 160) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")


def load_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def require_paths() -> None:
    paths = [
        GOBERNANZA_51,
        VALIDADOR_52_SCRIPT,
        VALIDADOR_52_XLSX,
        CONSOLIDADO_58,
        CONSOLIDADO_58_MANIFEST,
        MATRIZ_53,
        PRECARGA_5809,
    ]
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        raise FileNotFoundError("Faltan entradas obligatorias:\n" + "\n".join(missing))


def run_validator_52(output_dir: Path, timestamp: str) -> Dict[str, Any]:
    plan_path = output_dir / "plan_uso_validador_h52_paquete_validacion_humana_16_19.json"
    plan = {
        "descripcion": (
            "Paquete documental para validacion humana de 423 pendientes 16-19. "
            "No corrige, no recalcula 20-21, no genera carga ni SIES_READY."
        ),
        "columnas_usadas": [
            "CODIGO_UNICO",
            "PLAN_ESTUDIOS",
            "NUM_DOCUMENTO",
            "DV",
            "CODCLI_LISTA",
            "CODCLI",
            "ANO",
            "PERIODO",
            "CODRAMO",
            "ESTADO",
            "DESCRIPCION_ESTADO",
        ],
        "llaves_usadas": ["CODIGO_UNICO", "PLAN_ESTUDIOS", "NUM_DOCUMENTO", "DV", "CODCLI_LISTA"],
        "diccionarios_usados": ["DIC_ESTADO_PROMEDIOS"],
        "transformaciones": ["Preparacion documental de preguntas para validacion humana; sin correccion"],
        "campos_calculo": [],
        "anio_calculo": 2025,
        "genera_sies_ready": False,
        "modifica_fuentes_originales": False,
        "usa_todos_codcli_rut": False,
        "filtra_por_programa": True,
        "usa_inferencias_como_reglas_oficiales": False,
        "usa_estados_AEIR": True,
        "usa_codcli_lista": True,
        "mezcla_anual_acumulado": False,
        "modo": "PAQUETE_VALIDACION_HUMANA_NO_CORRECTIVO_16_19",
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    validator_timestamp = f"{timestamp}_H59"
    cmd = [
        sys.executable,
        str(VALIDADOR_52_SCRIPT),
        "--timestamp",
        validator_timestamp,
        "--plan-json",
        str(plan_path),
    ]
    completed = subprocess.run(cmd, check=False, capture_output=True, text=True)
    validator_dir = (
        REPO
        / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809"
        / f"VALIDADOR_GOBERNANZA_PRECALCULO_5809_{validator_timestamp}"
    )
    manifest_path = validator_dir / "manifest_validacion_gobernanza_precalculo_5809.json"
    manifest: Dict[str, Any] = {}
    if manifest_path.exists():
        manifest = load_json(manifest_path)
    result = {
        "plan_path": str(plan_path),
        "validator_timestamp": validator_timestamp,
        "validator_dir": str(validator_dir),
        "validator_manifest": str(manifest_path),
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "estado_validador": manifest.get("estado_validador", "SIN_MANIFEST"),
        "manifest": manifest,
    }
    if completed.returncode != 0:
        raise RuntimeError(
            "Validador hito 52 fallo.\nSTDOUT:\n"
            + completed.stdout
            + "\nSTDERR:\n"
            + completed.stderr
        )
    if result["estado_validador"] != ESTADO_VALIDADOR_REQUERIDO:
        raise RuntimeError(f"Validador hito 52 no apto: {result['estado_validador']}")
    return result


def read_precarga() -> pd.DataFrame:
    try:
        df = pd.read_csv(PRECARGA_5809, sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    except UnicodeDecodeError:
        df = pd.read_csv(PRECARGA_5809, sep=";", dtype=str, encoding="latin1").fillna("")
    df.insert(0, "FILA_5809", [str(i) for i in range(1, len(df) + 1)])
    return df


def read_matriz_values() -> pd.DataFrame:
    parts = []
    for sheet in [
        "03_DIF_SOLO_APROBADAS_19",
        "04_DIF_16_18_MULTIPLE",
        "05_SIN_REGISTROS_2025",
        "06_SIN_CODCLI_LISTA",
    ]:
        df = read_sheet(MATRIZ_53, sheet)
        df["HOJA_MATRIZ_53"] = sheet
        parts.append(df)
    matrix = pd.concat(parts, ignore_index=True)
    if matrix["FILA_5809"].duplicated().any():
        dup = matrix[matrix["FILA_5809"].duplicated(keep=False)]["FILA_5809"].tolist()
        raise RuntimeError("FILA_5809 duplicadas en matriz de valores: " + ", ".join(dup[:20]))
    return matrix


def value_pair(row: pd.Series, columns: str, suffix: str) -> str:
    mapping = {
        "16": ("CURSO_1ER_SEM_5809", "CURSO_1ER_SEM_RECALC", "CURSO_1ER_SEM"),
        "17": ("CURSO_2DO_SEM_5809", "CURSO_2DO_SEM_RECALC", "CURSO_2DO_SEM"),
        "18": ("UNIDADES_CURSADAS_5809", "UNIDADES_CURSADAS_RECALC", "UNIDADES_CURSADAS"),
        "19": ("UNIDADES_APROBADAS_5809", "UNIDADES_APROBADAS_RECALC_DICC", "UNIDADES_APROBADAS"),
    }
    out = []
    requested = [part.strip() for part in clean(columns).split(",") if part.strip()]
    if not requested:
        requested = ["16", "17", "18", "19"]
    for col in requested:
        current_field, recalc_field, label = mapping[col]
        field = current_field if suffix == "5809" else recalc_field
        out.append(f"{col} {label}={clean(row.get(field, ''))}")
    return " | ".join(out)


def question_for_case(row: pd.Series) -> str:
    route = clean(row.get("RUTA_RESOLUCION", ""))
    cause = clean(row.get("CAUSA_AUDITADA", ""))
    col = clean(row.get("COLUMNA_AFECTADA_16_19", ""))
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL" and cause == "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA":
        return "Para la columna 19 anual 2025, ¿deben contarse como aprobadas las actividades E=CONVALIDACION e I=HOMOLOGADO?"
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL" and cause == "PERIODO_NO_1_2":
        return "Para columnas 16-18, ¿como debe tratarse actividad 2025 con PERIODO distinto de 1 o 2?"
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL" and cause == "ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025":
        return "¿El estado academico observado permite cerrar el caso sin actividad 2025 y mantener los valores actuales de 16-19?"
    if route == "REQUIERE_MAPEO_INSTITUCIONAL":
        return "¿Cual es el CODCLI_LISTA institucional correcto para esta fila 5809 y este programa, o debe mantenerse sin mapeo?"
    if route == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA":
        return "¿Existe fuente academica 2025 complementaria para respaldar valores de 16-19, o se documenta ausencia de actividad?"
    if route == "REQUIERE_REVISION_FUNCIONAL":
        return f"¿Que valor validado debe mantenerse para las columnas {col} y cual es el fundamento funcional?"
    return "¿Se confirma que el caso debe permanecer bloqueado y sin correccion automatica?"


def expected_response_for_case(row: pd.Series) -> str:
    route = clean(row.get("RUTA_RESOLUCION", ""))
    cause = clean(row.get("CAUSA_AUDITADA", ""))
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL" and cause == "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA":
        return "SI/NO para contar E/I en columna 19, regla aplicable y responsable que valida."
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL":
        return "Decision funcional documentada, regla aplicable y valores autorizados para columnas afectadas."
    if route == "REQUIERE_MAPEO_INSTITUCIONAL":
        return "CODCLI_LISTA autorizado por fila/programa, o confirmacion de que no existe mapeo institucional valido."
    if route == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA":
        return "Fuente academica 2025, detalle de registros respaldatorios, o confirmacion documentada de ausencia."
    if route == "REQUIERE_REVISION_FUNCIONAL":
        return "Valor a mantener/corregir en 16-19, fundamento y autorizacion funcional."
    return "Confirmacion de bloqueo y motivo para no resolver automaticamente."


def observation_for_case(row: pd.Series) -> str:
    route = clean(row.get("RUTA_RESOLUCION", ""))
    if route == "RESOLUBLE_CON_DECISION_FUNCIONAL":
        return "Puede cerrarse solo con decision funcional explicita; no aplicar correccion en este paquete."
    if route == "REQUIERE_MAPEO_INSTITUCIONAL":
        return "No usar RUT ni N_CODCLI como sustituto de CODCLI_LISTA gobernado."
    if route == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA":
        return "No asumir cero ni actividad sin evidencia academica 2025 complementaria."
    return "Requiere revision puntual antes de cualquier accion operativa."


def build_package_cases(consolidated: pd.DataFrame, matrix: pd.DataFrame, precarga: pd.DataFrame) -> pd.DataFrame:
    matrix_by_fila = matrix.set_index("FILA_5809", drop=False)
    precarga_by_fila = precarga.set_index("FILA_5809", drop=False)
    rows: List[Dict[str, Any]] = []
    for _, row in consolidated.iterrows():
        fila = clean(row.get("FILA_5809", ""))
        mrow = matrix_by_fila.loc[fila] if fila in matrix_by_fila.index else pd.Series(dtype=object)
        prow = precarga_by_fila.loc[fila] if fila in precarga_by_fila.index else pd.Series(dtype=object)
        package_row = {
            "FILA_5809": fila,
            "NUM_DOCUMENTO": clean(prow.get("NUM_DOCUMENTO", "")),
            "DV": clean(prow.get("DV", "")),
            "CODIGO_UNICO": clean(row.get("CODIGO_UNICO", "")),
            "PLAN_ESTUDIOS": clean(row.get("PLAN_ESTUDIOS", "")),
            "CODCLI_LISTA": clean(row.get("CODCLI_LISTA_NORM", "")),
            "COLUMNA_AFECTADA_16_19": clean(row.get("COLUMNA_AFECTADA_16_19", "")),
            "VALOR_ACTUAL_5809": value_pair(mrow, clean(row.get("COLUMNA_AFECTADA_16_19", "")), "5809"),
            "VALOR_RECALCULADO_OBSERVADO": value_pair(mrow, clean(row.get("COLUMNA_AFECTADA_16_19", "")), "RECALC"),
            "CAUSA_AUDITADA": clean(row.get("CAUSA_AUDITADA", "")),
            "CAUSA_ORIGINAL_HITO_53": clean(row.get("CAUSA_ORIGINAL_HITO_53", "")),
            "RUTA_PRIMARIA": clean(row.get("RUTA_RESOLUCION", "")),
            "ACCION_SUGERIDA": clean(row.get("ACCION_SUGERIDA", "")),
            "PREGUNTA_CONCRETA_AREA_RESPONSABLE": question_for_case(row),
            "RESPUESTA_ESPERADA": expected_response_for_case(row),
            "PUEDE_CERRARSE_CON_DECISION_FUNCIONAL": "SI"
            if clean(row.get("RUTA_RESOLUCION", "")) == "RESOLUBLE_CON_DECISION_FUNCIONAL"
            else "NO",
            "REQUIERE_MAPEO_INSTITUCIONAL": "SI"
            if clean(row.get("RUTA_RESOLUCION", "")) == "REQUIERE_MAPEO_INSTITUCIONAL"
            else "NO",
            "REQUIERE_FUENTE_COMPLEMENTARIA": "SI"
            if clean(row.get("RUTA_RESOLUCION", "")) == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA"
            else "NO",
            "RIESGO": clean(row.get("RIESGO", "")),
            "PRIORIDAD": clean(row.get("PRIORIDAD", "")),
            "OBSERVACION": observation_for_case(row),
            "PUEDE_CORREGIRSE_AUTOMATICAMENTE": "NO",
            "REQUIERE_VALIDACION_HUMANA": "SI",
            "AFECTA_20_21": "NO EVALUADO / SEPARADO",
            "DECLARACION_CARGA": DECLARACION_CARGA,
        }
        rows.append(package_row)
    package = pd.DataFrame(rows)
    if len(package) != 423:
        raise RuntimeError(f"Paquete debe tener 423 casos; tiene {len(package)}")
    return package


def build_resumen_ejecutivo(package: pd.DataFrame) -> pd.DataFrame:
    rows = [
        {"Indicador": "Estado de carga", "Valor": DECLARACION_CARGA, "Detalle": "No se solicita carga ni archivo operativo."},
        {"Indicador": "Total casos pendientes 16-19", "Valor": len(package), "Detalle": "Universo proveniente del consolidado hito 58."},
        {
            "Indicador": "Decision funcional",
            "Valor": int((package["RUTA_PRIMARIA"] == "RESOLUBLE_CON_DECISION_FUNCIONAL").sum()),
            "Detalle": "Casos que requieren respuesta funcional explicita.",
        },
        {
            "Indicador": "Mapeo institucional",
            "Valor": int((package["RUTA_PRIMARIA"] == "REQUIERE_MAPEO_INSTITUCIONAL").sum()),
            "Detalle": "Casos que requieren CODCLI_LISTA/programa validado.",
        },
        {
            "Indicador": "Fuente complementaria",
            "Valor": int((package["RUTA_PRIMARIA"] == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA").sum()),
            "Detalle": "Casos sin evidencia 2025 suficiente en fuente actual.",
        },
        {
            "Indicador": "Revision funcional puntual",
            "Valor": int((package["RUTA_PRIMARIA"] == "REQUIERE_REVISION_FUNCIONAL").sum()),
            "Detalle": "Casos no explicados completamente por rutas anteriores.",
        },
        {
            "Indicador": "Correccion automatica ahora",
            "Valor": int((package["PUEDE_CORREGIRSE_AUTOMATICAMENTE"] == "SI").sum()),
            "Detalle": "Debe ser 0.",
        },
        {"Indicador": "20-21 acumulado", "Valor": "NO EVALUADO / SEPARADO", "Detalle": "Fuera de alcance de este paquete."},
    ]
    return pd.DataFrame(rows)


def build_question_catalog(package: pd.DataFrame) -> pd.DataFrame:
    return (
        package.groupby(["RUTA_PRIMARIA", "CAUSA_AUDITADA", "PREGUNTA_CONCRETA_AREA_RESPONSABLE", "RESPUESTA_ESPERADA"], dropna=False)
        .agg(
            CASOS=("FILA_5809", "count"),
            COLUMNAS_AFECTADAS=("COLUMNA_AFECTADA_16_19", lambda s: " | ".join(sorted(set(s)))),
            PRIORIDAD=("PRIORIDAD", "first"),
        )
        .reset_index()
        .sort_values(["PRIORIDAD", "RUTA_PRIMARIA", "CASOS"], ascending=[True, True, False])
    )


def build_required_fields() -> pd.DataFrame:
    fields = [
        ("FILA_5809", "Identificador fila 5809", "Obligatorio", "Numero de fila revisada."),
        ("RUTA_PRIMARIA", "Ruta propuesta", "Obligatorio", "No cambiar sin justificacion."),
        ("RESPUESTA_VALIDACION", "Respuesta humana", "Obligatorio", "SI/NO, CODCLI_LISTA, fuente o decision segun pregunta."),
        ("FUNDAMENTO", "Fundamento funcional o tecnico", "Obligatorio", "Regla, evidencia o motivo de rechazo."),
        ("RESPONSABLE", "Nombre/cargo responsable", "Obligatorio", "Persona o area que valida."),
        ("FECHA_RESPUESTA", "Fecha de respuesta", "Obligatorio", "YYYY-MM-DD."),
        ("VALOR_VALIDADO_16", "Valor validado columna 16", "Condicional", "Solo si aplica a la fila."),
        ("VALOR_VALIDADO_17", "Valor validado columna 17", "Condicional", "Solo si aplica a la fila."),
        ("VALOR_VALIDADO_18", "Valor validado columna 18", "Condicional", "Solo si aplica a la fila."),
        ("VALOR_VALIDADO_19", "Valor validado columna 19", "Condicional", "Solo si aplica a la fila."),
        ("CODCLI_LISTA_VALIDADO", "CODCLI_LISTA validado", "Condicional", "Obligatorio para ruta de mapeo institucional."),
        ("FUENTE_COMPLEMENTARIA", "Fuente academica adicional", "Condicional", "Obligatorio para ruta de fuente complementaria."),
        ("AUTORIZA_CORRECCION_FUTURA", "Autorizacion futura", "Obligatorio", "Solo habilita un paso posterior; no corrige este paquete."),
        ("OBSERVACION_RESPUESTA", "Observacion", "Opcional", "Aclaraciones del area responsable."),
    ]
    return pd.DataFrame(fields, columns=["Campo requerido", "Descripcion", "Obligatoriedad", "Uso esperado"])


def build_validation_df(validator: Dict[str, Any]) -> pd.DataFrame:
    rows = [
        {
            "Control": "Gobernanza integral hito 51",
            "Estado": "USADA",
            "Evidencia": str(GOBERNANZA_51),
            "Hash SHA256": sha256_file(GOBERNANZA_51),
            "Observacion": "Base obligatoria para fuentes, columnas, llaves y diccionarios.",
        },
        {
            "Control": "Validador hito 52",
            "Estado": validator["estado_validador"],
            "Evidencia": validator["validator_manifest"],
            "Hash SHA256": sha256_file(Path(validator["validator_manifest"])),
            "Observacion": "Compuerta ejecutada antes de generar hito 59.",
        },
        {
            "Control": "Consolidado hito 58",
            "Estado": "USADO_COMO_ENTRADA_PRINCIPAL",
            "Evidencia": str(CONSOLIDADO_58),
            "Hash SHA256": sha256_file(CONSOLIDADO_58),
            "Observacion": "Universo gobernado de 423 casos y rutas primarias.",
        },
        {
            "Control": "Manifest hito 58",
            "Estado": "USADO",
            "Evidencia": str(CONSOLIDADO_58_MANIFEST),
            "Hash SHA256": sha256_file(CONSOLIDADO_58_MANIFEST),
            "Observacion": "Confirma NO_LISTO_PARA_CARGA y no correccion automatica.",
        },
        {
            "Control": "Matriz hito 53",
            "Estado": "USADA_PARA_VALORES_OBSERVADOS",
            "Evidencia": str(MATRIZ_53),
            "Hash SHA256": sha256_file(MATRIZ_53),
            "Observacion": "Solo para valor actual 5809 y valor recalculado observado; no recalcula.",
        },
        {
            "Control": "Precarga oficial 5809",
            "Estado": "USADA_PARA_DOCUMENTO_DV",
            "Evidencia": str(PRECARGA_5809),
            "Hash SHA256": sha256_file(PRECARGA_5809),
            "Observacion": "Solo lectura para NUM_DOCUMENTO/DV; fuente no modificada.",
        },
        {
            "Control": "Plan validador hito 52 para paquete humano",
            "Estado": validator["estado_validador"],
            "Evidencia": validator["plan_path"],
            "Hash SHA256": sha256_file(Path(validator["plan_path"])),
            "Observacion": "Plan documental no correctivo.",
        },
    ]
    return pd.DataFrame(rows)


def build_dictamen(package: pd.DataFrame, validator: Dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Estado de carga": DECLARACION_CARGA,
                "Estado validador": validator["estado_validador"],
                "Fuentes originales modificadas": "NO",
                "Correcciones aplicadas": "NO",
                "SIES_READY generado": "NO",
                "Archivo de carga creado": "NO",
                "Recalculo 20-21": "NO",
                "Total casos paquete": len(package),
                "Correccion automatica habilitada": int((package["PUEDE_CORREGIRSE_AUTOMATICAMENTE"] == "SI").sum()),
                "Validacion humana requerida": int((package["REQUIERE_VALIDACION_HUMANA"] == "SI").sum()),
                "Dictamen global": (
                    "NO_LISTO_PARA_CARGA: paquete de validacion humana para 423 casos 16-19; "
                    "no aplica correccion automatica ni generacion de carga."
                ),
                "Proximo paso permitido": "Enviar expediente a responsables funcionales/institucionales y registrar respuestas.",
            }
        ]
    )


def build_fuentes_df(output_dir: Path, output_paths: Dict[str, Path], validator: Dict[str, Any]) -> pd.DataFrame:
    inputs = [
        ("INPUT", "Gobernanza integral hito 51", GOBERNANZA_51),
        ("INPUT", "Validador hito 52 xlsx base", VALIDADOR_52_XLSX),
        ("INPUT", "Validador hito 52 script", VALIDADOR_52_SCRIPT),
        ("INPUT", "Consolidado hito 58", CONSOLIDADO_58),
        ("INPUT", "Manifest consolidado hito 58", CONSOLIDADO_58_MANIFEST),
        ("INPUT", "Matriz hito 53", MATRIZ_53),
        ("INPUT", "Precarga oficial 5809", PRECARGA_5809),
        ("INPUT_DERIVADO", "Manifest validador hito 52 ejecucion hito 59", Path(validator["validator_manifest"])),
    ]
    rows = []
    for tipo, nombre, path in inputs:
        rows.append(
            {
                "Tipo": tipo,
                "Nombre": nombre,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() else "",
                "Uso": "Entrada o compuerta obligatoria; no modificada.",
            }
        )
    for key, path in output_paths.items():
        rows.append(
            {
                "Tipo": "OUTPUT",
                "Nombre": key,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() and key != "manifest" else "",
                "Uso": "Producto hito 59 no apto para carga. Manifest sin hash autorreferencial.",
            }
        )
    rows.append(
        {
            "Tipo": "OUTPUT_DIR",
            "Nombre": "Carpeta hito 59",
            "Ruta": str(output_dir),
            "Existe": "SI" if output_dir.exists() else "NO",
            "Hash SHA256": "",
            "Uso": "Contenedor de paquete de validacion humana.",
        }
    )
    return pd.DataFrame(rows)


def write_excel(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for logical_name, df in sheets.items():
            df.to_excel(writer, sheet_name=SHEETS.get(logical_name, logical_name)[:31], index=False)
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            values = [display(cell.value, 90) for cell in column_cells[:200]]
            width = min(max([len(v) for v in values] + [10]) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    wb.save(path)


def build_markdown(
    output_dir: Path,
    dictamen_df: pd.DataFrame,
    validation_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    questions_df: pd.DataFrame,
    manifest: Dict[str, Any],
) -> str:
    lines = [
        "# Informe Paquete Validacion Humana 16-19 5809",
        "",
        f"- Proceso: {PROCESO}",
        f"- Subproyecto: {SUBPROYECTO}",
        f"- Anio proceso: {ANIO_PROCESO}",
        f"- Anio referencia datos: {ANIO_REFERENCIA_DATOS}",
        f"- Declaracion carga: {DECLARACION_CARGA}",
        "- Fuentes originales modificadas: NO",
        "- Correcciones aplicadas: NO",
        "- SIES_READY generado: NO",
        "- Archivo de carga creado: NO",
        "- Recalculo 20-21: NO",
        f"- Carpeta: {output_dir}",
        "",
        "## Dictamen",
        "",
        markdown_table(dictamen_df),
        "",
        "## Validacion de gobernanza",
        "",
        markdown_table(validation_df[["Control", "Estado", "Observacion"]]),
        "",
        "## Resumen ejecutivo",
        "",
        markdown_table(resumen_df),
        "",
        "## Preguntas principales",
        "",
        markdown_table(
            questions_df[
                [
                    "RUTA_PRIMARIA",
                    "CAUSA_AUDITADA",
                    "CASOS",
                    "PREGUNTA_CONCRETA_AREA_RESPONSABLE",
                    "RESPUESTA_ESPERADA",
                ]
            ],
            max_rows=20,
        ),
        "",
        "## Criterio operativo",
        "",
        "- Este paquete solo prepara validacion humana.",
        "- No solicita carga SIES y no habilita correcciones automaticas.",
        "- Los campos 20-21 acumulados quedan separados y no evaluados.",
        "- Toda respuesta debe registrarse antes de cualquier paso posterior.",
        "",
        "## Archivos",
        "",
        f"- Excel: {manifest.get('archivos', {}).get('excel', '')}",
        f"- Informe: {manifest.get('archivos', {}).get('informe', '')}",
        f"- Manifest: {manifest.get('archivos', {}).get('manifest', '')}",
        f"- Script: {manifest.get('archivos', {}).get('script', '')}",
        f"- Borrador correo: {manifest.get('archivos', {}).get('correo', '')}",
        "",
    ]
    return "\n".join(lines)


def build_email_draft(output_paths: Dict[str, Path]) -> str:
    return "\n".join(
        [
            "# Borrador correo validacion 16-19 5809",
            "",
            "Asunto: Validacion humana requerida para 423 casos pendientes 16-19 - Avance Curricular 5809",
            "",
            "Estimadas/os,",
            "",
            "Compartimos el paquete de validacion humana para los casos pendientes de columnas 16-19 del archivo 5809 Matrícula Avance Curricular.",
            "",
            "Este envio no solicita carga SIES, no contiene archivo de carga y no aplica correcciones automaticas. Las fuentes originales no fueron modificadas.",
            "",
            "Resumen del paquete:",
            "",
            "- Total de casos pendientes: 423.",
            "- 256 casos requieren decision funcional.",
            "- 102 casos requieren mapeo institucional.",
            "- 57 casos requieren fuente academica complementaria.",
            "- 8 casos requieren revision funcional puntual.",
            "- Las columnas 20-21 de acumulado estan separadas y no fueron evaluadas en este paquete.",
            "",
            "Solicitamos revisar las hojas correspondientes a cada ruta y responder las preguntas indicadas en la hoja 07_PREGUNTAS_PARA_VALIDAR. La respuesta debe incluir fundamento, responsable, fecha y, cuando corresponda, valor validado o CODCLI_LISTA/fuente complementaria.",
            "",
            "Archivos de referencia del paquete:",
            "",
            f"- Excel: {output_paths['excel']}",
            f"- Informe: {output_paths['informe']}",
            f"- Manifest: {output_paths['manifest']}",
            "",
            "Declaracion de estado: NO_LISTO_PARA_CARGA.",
            "",
            "Saludos,",
            "Equipo Avance Curricular SIES 2026",
            "",
        ]
    )


def build_manifest(
    timestamp: str,
    output_dir: Path,
    output_paths: Dict[str, Path],
    validator: Dict[str, Any],
    package: pd.DataFrame,
    sheets: Dict[str, pd.DataFrame],
) -> Dict[str, Any]:
    return {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": timestamp,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "correcciones_aplicadas": "NO",
        "sies_ready_generado": "NO",
        "archivo_carga_creado": "NO",
        "recalculo_20_21": "NO",
        "estado_validador_hito_52": validator["estado_validador"],
        "carpeta_salida": str(output_dir),
        "archivos": {key: str(path) for key, path in output_paths.items()},
        "entradas_obligatorias": {
            "gobernanza_51": str(GOBERNANZA_51),
            "validador_52_script": str(VALIDADOR_52_SCRIPT),
            "validador_52_xlsx": str(VALIDADOR_52_XLSX),
            "consolidado_58": str(CONSOLIDADO_58),
            "manifest_consolidado_58": str(CONSOLIDADO_58_MANIFEST),
            "matriz_53_valores_observados": str(MATRIZ_53),
            "precarga_5809_documento_dv": str(PRECARGA_5809),
        },
        "validador_hito_52_ejecucion": {
            "plan": validator["plan_path"],
            "manifest": validator["validator_manifest"],
            "carpeta": validator["validator_dir"],
            "returncode": validator["returncode"],
        },
        "conteos": {
            "casos_paquete": len(package),
            "decision_funcional": int((package["RUTA_PRIMARIA"] == "RESOLUBLE_CON_DECISION_FUNCIONAL").sum()),
            "mapeo_institucional": int((package["RUTA_PRIMARIA"] == "REQUIERE_MAPEO_INSTITUCIONAL").sum()),
            "fuente_complementaria": int((package["RUTA_PRIMARIA"] == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA").sum()),
            "revision_funcional": int((package["RUTA_PRIMARIA"] == "REQUIERE_REVISION_FUNCIONAL").sum()),
            "correccion_automatica": int((package["PUEDE_CORREGIRSE_AUTOMATICAMENTE"] == "SI").sum()),
            "hojas_excel": {name: len(df) for name, df in sheets.items()},
        },
        "hashes_salida": {
            key: sha256_file(path) for key, path in output_paths.items() if path.exists() and key != "manifest"
        },
    }


def print_console(
    dictamen_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    output_paths: Dict[str, Path],
    output_dir: Path,
    validator_state: str,
) -> None:
    print("PAQUETE VALIDACION HUMANA 16-19 5809 — GENERADO")
    print("Fuentes originales modificadas: NO")
    print("Declaracion carga: NO_LISTO_PARA_CARGA")
    print(f"Estado validador: {validator_state}")
    print()
    print("DICTAMEN")
    print(dictamen_df.to_string(index=False))
    print()
    print("RESUMEN EJECUTIVO")
    print(resumen_df.to_string(index=False))
    print()
    print("ARCHIVOS")
    print(f"Excel: {output_paths['excel']}")
    print(f"Informe: {output_paths['informe']}")
    print(f"Manifest: {output_paths['manifest']}")
    print(f"Script: {output_paths['script']}")
    print(f"Borrador correo: {output_paths['correo']}")
    print(f"Carpeta hito: {output_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Paquete hito 59 validacion humana 16-19 5809")
    parser.add_argument("--timestamp", default=None, help="Timestamp YYYYMMDD_HHMMSS. Si se omite, usa fecha/hora actual.")
    args = parser.parse_args()

    timestamp = args.timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = BASE_59 / f"PAQUETE_VALIDACION_HUMANA_16_19_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths = {
        "excel": output_dir / OUTPUT_FILES["excel"],
        "informe": output_dir / OUTPUT_FILES["informe"],
        "manifest": output_dir / OUTPUT_FILES["manifest"],
        "script": output_dir / OUTPUT_FILES["script"],
        "correo": output_dir / OUTPUT_FILES["correo"],
    }

    require_paths()
    current_script = Path(__file__).resolve()
    target_script = output_paths["script"].resolve()
    if current_script != target_script:
        shutil.copy2(current_script, target_script)

    validator = run_validator_52(output_dir, timestamp)
    h58_manifest = load_json(CONSOLIDADO_58_MANIFEST)
    if h58_manifest.get("declaracion_carga") != DECLARACION_CARGA:
        raise RuntimeError("Consolidado hito 58 no declara NO_LISTO_PARA_CARGA")
    consolidated = read_sheet(CONSOLIDADO_58, "03_CONSOLIDADO_423")
    matrix = read_matriz_values()
    precarga = read_precarga()
    package = build_package_cases(consolidated, matrix, precarga)

    expected_counts = {
        "RESOLUBLE_CON_DECISION_FUNCIONAL": 256,
        "REQUIERE_MAPEO_INSTITUCIONAL": 102,
        "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA": 57,
        "REQUIERE_REVISION_FUNCIONAL": 8,
    }
    observed_counts = package["RUTA_PRIMARIA"].value_counts().to_dict()
    for route, expected in expected_counts.items():
        if int(observed_counts.get(route, 0)) != expected:
            raise RuntimeError(f"Conteo ruta {route} esperado {expected}, observado {observed_counts.get(route, 0)}")

    dictamen_df = build_dictamen(package, validator)
    validation_df = build_validation_df(validator)
    resumen_df = build_resumen_ejecutivo(package)
    questions_df = build_question_catalog(package)
    required_fields_df = build_required_fields()
    priority_df = package.sort_values(["PRIORIDAD", "RUTA_PRIMARIA", "FILA_5809"]).copy()
    no_auto_df = package[
        [
            "FILA_5809",
            "NUM_DOCUMENTO",
            "DV",
            "CODIGO_UNICO",
            "PLAN_ESTUDIOS",
            "RUTA_PRIMARIA",
            "CAUSA_AUDITADA",
            "PUEDE_CORREGIRSE_AUTOMATICAMENTE",
            "OBSERVACION",
            "DECLARACION_CARGA",
        ]
    ].copy()

    sheets = OrderedDict(
        [
            ("00_DICTAMEN", dictamen_df),
            ("01_VALIDACION_GOBERNANZA", validation_df),
            ("02_RESUMEN_EJECUTIVO", resumen_df),
            (
                "03_DECISION_FUNCIONAL_256",
                package[package["RUTA_PRIMARIA"] == "RESOLUBLE_CON_DECISION_FUNCIONAL"],
            ),
            (
                "04_MAPEO_INSTITUCIONAL_102",
                package[package["RUTA_PRIMARIA"] == "REQUIERE_MAPEO_INSTITUCIONAL"],
            ),
            (
                "05_FUENTE_COMPLEMENTARIA_57",
                package[package["RUTA_PRIMARIA"] == "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA"],
            ),
            (
                "06_REVISION_FUNCIONAL_8",
                package[package["RUTA_PRIMARIA"] == "REQUIERE_REVISION_FUNCIONAL"],
            ),
            ("07_PREGUNTAS_PARA_VALIDAR", questions_df),
            ("08_CAMPOS_REQUERIDOS_RESPUESTA", required_fields_df),
            ("09_PRIORIDAD_REVISION", priority_df),
            ("10_NO_CORREGIR_AUTOMATICO", no_auto_df),
        ]
    )
    sheets["11_FUENTES"] = build_fuentes_df(output_dir, output_paths, validator)
    write_excel(output_paths["excel"], sheets)

    manifest = build_manifest(timestamp, output_dir, output_paths, validator, package, sheets)
    output_paths["informe"].write_text(
        build_markdown(output_dir, dictamen_df, validation_df, resumen_df, questions_df, manifest),
        encoding="utf-8",
    )
    output_paths["correo"].write_text(build_email_draft(output_paths), encoding="utf-8")
    sheets["11_FUENTES"] = build_fuentes_df(output_dir, output_paths, validator)
    write_excel(output_paths["excel"], sheets)
    manifest = build_manifest(timestamp, output_dir, output_paths, validator, package, sheets)
    output_paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print_console(dictamen_df, resumen_df, output_paths, output_dir, validator["estado_validador"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
