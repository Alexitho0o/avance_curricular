from pathlib import Path
from datetime import datetime
import pandas as pd
import numpy as np
import json
import hashlib
import shutil
import zipfile
import re

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =============================================================================
# CONTEXTO
# =============================================================================

PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Aplicación regla gobernada sin actividad académica 5809"
ANIO_REFERENCIA = 2025

RAIZ = Path("/Users/alexi/Documents/GitHub/avance_curricular").resolve()
DESKTOP = Path.home() / "Desktop"

BASE_30 = (
    RAIZ
    / "avance_curricular_2026"
    / "30_mapeo_resolucion_bloqueos_5809_2025"
)

AUDITORIA_CORREGIDA = (
    RAIZ
    / "avance_curricular_2026"
    / "29_auditoria_corregida_logica_2025_5809"
    / "AUDITORIA_CORREGIDA_LOGICA_2025_5809.xlsx"
)

PROMEDIOS = (
    RAIZ
    / "avance_curricular_2026"
    / "25_actualizacion_fuente_promedios"
    / "CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059"
    / "00_FUENTE_CONGELADA"
    / "PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
)

CSV_5809 = (
    RAIZ
    / "avance_curricular_2026"
    / "14_sies_ready_5809"
    / "SIES_READY_TECNICO_5809_22_COLUMNAS_20260704_151418"
    / "01_ARCHIVO_SIES_READY_TECNICO"
    / "5809_MATRICULA_AVANCE_CURRICULAR_2026_SIES_READY_TECNICO.csv"
)

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

SALIDA = (
    RAIZ
    / "avance_curricular_2026"
    / "31_aplicacion_regla_gobernada_sin_actividad_5809_2025"
    / f"REGLA_GOBERNADA_SIN_ACTIVIDAD_5809_2025_{timestamp}"
)

RESULTADOS = SALIDA / "02_RESULTADOS"
AUDITORIA = SALIDA / "03_AUDITORIA"
REPORTES = SALIDA / "04_REPORTES"
MANIFEST_DIR = SALIDA / "05_MANIFEST"

for carpeta in [RESULTADOS, AUDITORIA, REPORTES, MANIFEST_DIR]:
    carpeta.mkdir(parents=True, exist_ok=True)

CARPETA_ESCRITORIO = (
    DESKTOP
    / f"AVANCE_CURRICULAR_2026_REGLA_GOBERNADA_SIN_ACTIVIDAD_5809_2025_{timestamp}"
)

CARPETA_ESCRITORIO.mkdir(parents=True, exist_ok=True)

COLUMNAS_5809 = [
    "CODIGO_IES_NUM",
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
    "VIGENCIA",
]

ESTADOS_GOBERNADOS_SIN_ACTIVIDAD = {
    "ELIMINADO",
    "ELIMINADA",
    "ELIMINADO SIN ASIGNATURAS",
    "ELIMINADO_SIN_ASIGNATURAS",
    "VIGENTE SIN ACTIVIDAD ACADEMICA",
    "VIGENTE SIN ACTIVIDAD ACADÉMICA",
    "VIGENTE_SIN_ACTIVIDAD_ACADEMICA",
    "SIN ACTIVIDAD ACADEMICA",
    "SIN ACTIVIDAD ACADÉMICA",
}


# =============================================================================
# UTILIDADES
# =============================================================================

def sha256(path):
    if not path.exists():
        return "NO_EXISTE"

    h = hashlib.sha256()

    with path.open("rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)

    return h.hexdigest()


def json_safe(obj):
    if isinstance(obj, Path):
        return str(obj)

    if isinstance(obj, np.integer):
        return int(obj)

    if isinstance(obj, np.floating):
        return float(obj)

    if isinstance(obj, np.bool_):
        return bool(obj)

    return obj


def dumps_json_safe(obj):
    return json.dumps(
        obj,
        ensure_ascii=False,
        indent=2,
        default=json_safe,
    )


def norm_text(x):
    txt = str(x or "").strip().upper()

    reemplazos = {
        "Á": "A",
        "É": "E",
        "Í": "I",
        "Ó": "O",
        "Ú": "U",
        "Ñ": "N",
    }

    for a, b in reemplazos.items():
        txt = txt.replace(a, b)

    txt = re.sub(r"\s+", " ", txt)
    return txt.strip()


def norm_col(x):
    txt = norm_text(x)
    txt = re.sub(r"[^A-Z0-9]+", "_", txt)
    txt = re.sub(r"_+", "_", txt)
    return txt.strip("_")


def only_digits(x):
    return re.sub(r"\D", "", str(x or ""))


def normalizar_rut_libre(x):
    txt = str(x or "").strip().upper()
    txt = txt.replace(".", "").replace(" ", "")

    if "-" in txt:
        cuerpo, dv = txt.split("-", 1)
        cuerpo = only_digits(cuerpo)
        dv = re.sub(r"[^0-9K]", "", dv)

        if cuerpo and dv:
            return f"{cuerpo}-{dv[-1]}"

        return ""

    limpio = re.sub(r"[^0-9K]", "", txt)

    if len(limpio) >= 2:
        return f"{limpio[:-1]}-{limpio[-1]}"

    return ""


def normalizar_rut_num_dv(num, dv):
    n = only_digits(num)
    d = re.sub(r"[^0-9K]", "", str(dv or "").upper())

    if n and d:
        return f"{n}-{d[-1]}"

    return ""


def to_bool(x):
    if isinstance(x, bool):
        return x

    txt = norm_text(x)

    return txt in ["TRUE", "VERDADERO", "SI", "SÍ", "1", "OK"]


def to_int(x):
    try:
        return int(float(str(x).replace(",", ".")))
    except Exception:
        return 0


def read_excel(path, sheet):
    try:
        return pd.read_excel(
            path,
            sheet_name=sheet,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl",
        )
    except Exception:
        return pd.DataFrame()


def detectar_columna(df, candidatos_exactos=None, contiene=None, excluir=None):
    candidatos_exactos = candidatos_exactos or []
    contiene = contiene or []
    excluir = excluir or []

    if df.empty:
        return None

    mapa = {norm_col(c): c for c in df.columns}

    for cand in candidatos_exactos:
        key = norm_col(cand)

        if key in mapa:
            return mapa[key]

    for c in df.columns:
        nc = norm_col(c)

        if all(token in nc for token in contiene) and not any(token in nc for token in excluir):
            return c

    return None


def detectar_col_codcli(df):
    return detectar_columna(
        df,
        candidatos_exactos=[
            "CODCLI",
            "COD_CLI",
            "CODIGO_CLIENTE",
            "CODIGO_ALUMNO",
            "CODALUMNO",
        ],
    )


def detectar_col_rut(df):
    return detectar_columna(
        df,
        candidatos_exactos=[
            "RUT",
            "RUN",
            "NUM_DOCUMENTO",
            "NUMDOCUMENTO",
            "DOCUMENTO",
            "RUT_ALUMNO",
        ],
    ) or detectar_columna(df, contiene=["RUT"])


def detectar_col_dv(df):
    return detectar_columna(
        df,
        candidatos_exactos=[
            "DV",
            "DVRUT",
            "DIGITO_VERIFICADOR",
            "DIG_VERIFICADOR",
        ],
    ) or detectar_columna(df, contiene=["DV"])


def detectar_col_estado_academico(df, nombre_hoja):
    """
    Evita usar Hoja1/ESTADO de asignaturas como estado académico del estudiante.
    Busca estados institucionales en hojas tipo DatosAlumnos / Matrícula.
    """

    if df.empty:
        return None

    hoja_norm = norm_col(nombre_hoja)

    candidatos_fuertes = [
        "ESTADO_ACADEMICO",
        "ESTADO ACADÉMICO",
        "ESTADO ACADEMICO",
        "ESTADO_ALUMNO",
        "ESTADO ALUMNO",
        "ESTADO_MATRICULA",
        "ESTADO MATRÍCULA",
        "ESTADO MATRICULA",
        "SITUACION_ACADEMICA",
        "SITUACIÓN ACADÉMICA",
        "SITUACION ACADEMICA",
        "SITUACION_ALUMNO",
        "SITUACIÓN ALUMNO",
        "CONDICION_ACADEMICA",
        "CONDICIÓN ACADÉMICA",
        "CONDICION_ALUMNO",
    ]

    col = detectar_columna(df, candidatos_exactos=candidatos_fuertes)

    if col:
        return col

    if hoja_norm in ["HOJA1"]:
        return None

    for c in df.columns:
        nc = norm_col(c)

        if "ESTADO" in nc and not any(ex in nc for ex in ["RAMO", "ASIGNATURA", "CURSO", "NOTA"]):
            return c

        if "SITUACION" in nc and not any(ex in nc for ex in ["RAMO", "ASIGNATURA", "CURSO", "NOTA"]):
            return c

        if "CONDICION" in nc and not any(ex in nc for ex in ["RAMO", "ASIGNATURA", "CURSO", "NOTA"]):
            return c

    return None


def es_estado_gobernado_sin_actividad(x):
    txt = norm_text(x)

    if txt in {norm_text(e) for e in ESTADOS_GOBERNADOS_SIN_ACTIVIDAD}:
        return True

    if "ELIMINAD" in txt:
        return True

    if "SIN ACTIVIDAD" in txt and "ACADEMIC" in txt:
        return True

    return False


def validar_valores_sin_actividad(row):
    return (
        norm_text(row.get("CURSO_1ER_SEM", "")) == "NO"
        and norm_text(row.get("CURSO_2DO_SEM", "")) == "NO"
        and to_int(row.get("UNIDADES_CURSADAS", 0)) == 0
        and to_int(row.get("UNIDADES_APROBADAS", 0)) == 0
        and to_int(row.get("UNID_CURSADAS_TOTAL", 0)) == 0
        and to_int(row.get("UNID_APROBADAS_TOTAL", 0)) == 0
    )


def validar_xlsx(path):
    with zipfile.ZipFile(path, "r") as zf:
        corrupto = zf.testzip()

        if corrupto is not None:
            raise RuntimeError(f"XLSX corrupto en {corrupto}")

    wb = load_workbook(path, read_only=True)
    hojas = wb.sheetnames
    wb.close()

    return hojas


def aplicar_formato(path):
    wb = load_workbook(path)

    fill_header = PatternFill("solid", fgColor="1F4E78")
    fill_ok = PatternFill("solid", fgColor="C6EFCE")
    fill_warn = PatternFill("solid", fgColor="FFEB9C")
    fill_bad = PatternFill("solid", fgColor="FFC7CE")
    fill_info = PatternFill("solid", fgColor="D9EAF7")

    font_header = Font(color="FFFFFF", bold=True)
    font_ok = Font(color="006100", bold=True)
    font_warn = Font(color="9C6500", bold=True)
    font_bad = Font(color="9C0006", bold=True)

    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for ws in wb.worksheets:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = align
                cell.border = border

        ws.freeze_panes = "A2"

        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, 2500),
            min_col=1,
            max_col=ws.max_column,
        ):
            for cell in row:
                cell.alignment = align
                cell.border = border

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0

            for row_idx in range(1, min(ws.max_row, 800) + 1):
                val = ws.cell(row_idx, col_idx).value

                if val is not None:
                    max_len = max(max_len, len(str(val)))

            header = str(ws.cell(1, col_idx).value or "").upper()

            width = min(max(max_len + 2, 12), 70)

            if any(k in header for k in ["RUTA", "ACCION", "MOTIVO", "OBSERVACION", "DETALLE", "CRITERIO", "REGLA"]):
                width = 70

            ws.column_dimensions[col_letter].width = width

        for row_idx in range(2, min(ws.max_row, 2500) + 1):
            texto_fila = " ".join(
                str(ws.cell(row_idx, c).value or "").upper()
                for c in range(1, ws.max_column + 1)
            )

            if "BLOQUEO" in texto_fila or "NO_CALZA" in texto_fila:
                fill = fill_bad
                font = font_bad
            elif "REVISAR" in texto_fila or "PENDIENTE" in texto_fila or "DIFERENCIA" in texto_fila:
                fill = fill_warn
                font = font_warn
            elif "OK" in texto_fila or "EXCEPCION_GOBERNADA" in texto_fila:
                fill = fill_ok
                font = font_ok
            else:
                fill = None
                font = None

            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row_idx, c).fill = fill
                    ws.cell(row_idx, c).font = font

        if ws.title.startswith("00") or "DICTAMEN" in ws.title:
            ws.sheet_properties.tabColor = "5B9BD5"
        elif "BLOQUEO" in ws.title or "PENDIENTE" in ws.title:
            ws.sheet_properties.tabColor = "FFC000"
        elif "GOBERNADA" in ws.title or "RESUELT" in ws.title:
            ws.sheet_properties.tabColor = "70AD47"
        else:
            ws.sheet_properties.tabColor = "A6A6A6"

    wb.save(path)


# =============================================================================
# ENTRADAS
# =============================================================================

bloqueos = []

if not BASE_30.exists():
    bloqueos.append(f"No existe carpeta base mapeo 30: {BASE_30}")

if not AUDITORIA_CORREGIDA.exists():
    bloqueos.append(f"No existe auditoría corregida: {AUDITORIA_CORREGIDA}")

if not PROMEDIOS.exists():
    bloqueos.append(f"No existe PROMEDIOS: {PROMEDIOS}")

if not CSV_5809.exists():
    bloqueos.append(f"No existe CSV 5809: {CSV_5809}")

if bloqueos:
    reporte = REPORTES / "BLOQUEO_ENTRADAS.md"
    reporte.write_text(
        "# Bloqueo de entradas\n\n"
        + "\n".join(f"- {b}" for b in bloqueos)
        + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 120)
    print("BLOQUEO DE ENTRADAS")
    print("=" * 120)

    for b in bloqueos:
        print(f"- {b}")

    print(f"Reporte: {reporte}")
    print("=" * 120)
    raise SystemExit(1)


carpetas_30 = sorted(
    [p for p in BASE_30.glob("MAPEO_RESOLUCION_BLOQUEOS_5809_2025_*") if p.is_dir()],
    key=lambda p: p.stat().st_mtime,
    reverse=True,
)

if not carpetas_30:
    raise SystemExit(f"BLOQUEO: no hay salidas del paso 30 en {BASE_30}")

ULTIMA_30 = carpetas_30[0]

EXCEL_MAPEO_30 = ULTIMA_30 / "02_RESULTADOS" / "MAPEO_RESOLUCION_BLOQUEOS_5809_LOGICA_2025.xlsx"

if not EXCEL_MAPEO_30.exists():
    raise SystemExit(f"BLOQUEO: no existe Excel del paso 30: {EXCEL_MAPEO_30}")


# =============================================================================
# CARGA BASES
# =============================================================================

df_5809 = pd.read_csv(
    CSV_5809,
    sep=";",
    header=None,
    names=COLUMNAS_5809,
    dtype=str,
    encoding="cp1252",
    keep_default_na=False,
)

df_5809["RUT_NORMALIZADO"] = df_5809.apply(
    lambda r: normalizar_rut_num_dv(r["NUM_DOCUMENTO"], r["DV"]),
    axis=1,
)

for col in [
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
]:
    df_5809[col] = pd.to_numeric(df_5809[col], errors="coerce").fillna(0).astype(int)


aud_anual = read_excel(AUDITORIA_CORREGIDA, "AUDITORIA_ANUAL_2025_FILA_A_FILA")
filas_revisar_anual = read_excel(AUDITORIA_CORREGIDA, "FILAS_REVISAR_ANUAL_2025")
aud_acum = read_excel(AUDITORIA_CORREGIDA, "AUDITORIA_ACUMULADO_2025_FILA_A_FILA")
catalogo_estado = read_excel(AUDITORIA_CORREGIDA, "CATALOGO_ESTADO")
mapeo_prev = read_excel(AUDITORIA_CORREGIDA, "MAPEO_CODCLI_RUT")

mapeo_30_p1 = read_excel(EXCEL_MAPEO_30, "P1_SIN_CODCLI_CANDIDATOS")
mapeo_30_p2 = read_excel(EXCEL_MAPEO_30, "P2_DIFERENCIAS_16_19")
mapeo_30_p3 = read_excel(EXCEL_MAPEO_30, "P3_ESTADOS_NO_CLASIFICADOS")
mapeo_30_p4 = read_excel(EXCEL_MAPEO_30, "P4_ACUMULADO_REVISAR")

if aud_anual.empty:
    raise SystemExit("BLOQUEO: no se pudo leer AUDITORIA_ANUAL_2025_FILA_A_FILA.")

if aud_acum.empty:
    raise SystemExit("BLOQUEO: no se pudo leer AUDITORIA_ACUMULADO_2025_FILA_A_FILA.")


# =============================================================================
# MAPA ESTADO ACADÉMICO INSTITUCIONAL
# =============================================================================

xls_prom = pd.ExcelFile(PROMEDIOS, engine="openpyxl")

mapeo_estado = []
diagnostico_estado_cols = []

for hoja in xls_prom.sheet_names:
    df = pd.read_excel(
        PROMEDIOS,
        sheet_name=hoja,
        dtype=str,
        keep_default_na=False,
        engine="openpyxl",
    )

    col_codcli = detectar_col_codcli(df)
    col_rut = detectar_col_rut(df)
    col_dv = detectar_col_dv(df)
    col_estado = detectar_col_estado_academico(df, hoja)

    diagnostico_estado_cols.append({
        "HOJA": hoja,
        "FILAS": len(df),
        "COLUMNAS": len(df.columns),
        "CODCLI": col_codcli or "",
        "RUT": col_rut or "",
        "DV": col_dv or "",
        "ESTADO_ACADEMICO_DETECTADO": col_estado or "",
        "USADA_PARA_ESTADO_ACADEMICO": "SI" if col_estado and (col_codcli or col_rut) else "NO",
    })

    if not col_estado:
        continue

    if not col_codcli and not col_rut:
        continue

    temp = pd.DataFrame()
    temp["HOJA_ORIGEN_ESTADO"] = hoja

    if col_codcli:
        temp["CODCLI"] = df[col_codcli].astype(str).str.strip()
    else:
        temp["CODCLI"] = ""

    if col_rut and col_dv and col_rut != col_dv:
        temp["RUT_NORMALIZADO"] = df.apply(
            lambda r: normalizar_rut_num_dv(r[col_rut], r[col_dv]),
            axis=1,
        )
    elif col_rut:
        temp["RUT_NORMALIZADO"] = df[col_rut].apply(normalizar_rut_libre)
    else:
        temp["RUT_NORMALIZADO"] = ""

    temp["ESTADO_ACADEMICO_ORIGINAL"] = df[col_estado].astype(str).str.strip()
    temp["ESTADO_ACADEMICO_NORM"] = temp["ESTADO_ACADEMICO_ORIGINAL"].apply(norm_text)
    temp["ES_REGLA_GOBERNADA_SIN_ACTIVIDAD"] = temp["ESTADO_ACADEMICO_ORIGINAL"].apply(es_estado_gobernado_sin_actividad)

    temp = temp[
        (temp["CODCLI"].ne("") | temp["RUT_NORMALIZADO"].ne(""))
        & temp["ESTADO_ACADEMICO_ORIGINAL"].ne("")
    ].copy()

    mapeo_estado.append(temp)

diagnostico_estado_cols_df = pd.DataFrame(diagnostico_estado_cols)

if mapeo_estado:
    mapa_estado_total = pd.concat(mapeo_estado, ignore_index=True).drop_duplicates()
else:
    mapa_estado_total = pd.DataFrame(
        columns=[
            "HOJA_ORIGEN_ESTADO",
            "CODCLI",
            "RUT_NORMALIZADO",
            "ESTADO_ACADEMICO_ORIGINAL",
            "ESTADO_ACADEMICO_NORM",
            "ES_REGLA_GOBERNADA_SIN_ACTIVIDAD",
        ]
    )

estado_por_rut = (
    mapa_estado_total[mapa_estado_total["RUT_NORMALIZADO"].ne("")]
    .groupby("RUT_NORMALIZADO", dropna=False)
    .agg(
        ESTADOS_ACADEMICOS=("ESTADO_ACADEMICO_ORIGINAL", lambda s: " | ".join(sorted(set(map(str, s))))),
        HOJAS_ESTADO=("HOJA_ORIGEN_ESTADO", lambda s: " | ".join(sorted(set(map(str, s))))),
        TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD=("ES_REGLA_GOBERNADA_SIN_ACTIVIDAD", "max"),
    )
    .reset_index()
)

estado_por_codcli = (
    mapa_estado_total[mapa_estado_total["CODCLI"].ne("")]
    .groupby("CODCLI", dropna=False)
    .agg(
        ESTADOS_ACADEMICOS_CODCLI=("ESTADO_ACADEMICO_ORIGINAL", lambda s: " | ".join(sorted(set(map(str, s))))),
        HOJAS_ESTADO_CODCLI=("HOJA_ORIGEN_ESTADO", lambda s: " | ".join(sorted(set(map(str, s))))),
        TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI=("ES_REGLA_GOBERNADA_SIN_ACTIVIDAD", "max"),
    )
    .reset_index()
)


# =============================================================================
# ENRIQUECER AUDITORÍA ANUAL CON ESTADO ACADÉMICO
# =============================================================================

base = aud_anual.copy()

if "RUT_NORMALIZADO" not in base.columns:
    raise SystemExit("BLOQUEO: AUDITORIA_ANUAL no tiene RUT_NORMALIZADO.")

for col in [
    "CODCLI_LISTA",
    "CODCLI_CANTIDAD",
    "DICTAMEN_ANUAL_2025",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
]:
    if col not in base.columns:
        base[col] = ""

base = base.merge(
    estado_por_rut,
    on="RUT_NORMALIZADO",
    how="left",
)

base["ESTADOS_ACADEMICOS"] = base["ESTADOS_ACADEMICOS"].fillna("")
base["HOJAS_ESTADO"] = base["HOJAS_ESTADO"].fillna("")
base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD"] = base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD"].fillna(False).astype(bool)

# Reforzar por CODCLI cuando el RUT no haya traído estado.
filas_codcli_estado = []

for idx, row in base.iterrows():
    lista = str(row.get("CODCLI_LISTA", "")).split("|")
    lista = [x.strip() for x in lista if x.strip()]

    estados = []
    hojas = []
    flag = False

    for cod in lista:
        temp = estado_por_codcli[estado_por_codcli["CODCLI"].eq(cod)]

        if temp.empty:
            continue

        for _, e in temp.iterrows():
            estados.append(str(e.get("ESTADOS_ACADEMICOS_CODCLI", "")))
            hojas.append(str(e.get("HOJAS_ESTADO_CODCLI", "")))
            flag = flag or bool(e.get("TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI", False))

    filas_codcli_estado.append({
        "_IDX": idx,
        "ESTADOS_ACADEMICOS_DESDE_CODCLI": " | ".join(sorted(set([x for x in estados if x]))),
        "HOJAS_ESTADO_DESDE_CODCLI": " | ".join(sorted(set([x for x in hojas if x]))),
        "TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI": flag,
    })

codcli_estado_df = pd.DataFrame(filas_codcli_estado)

base = base.merge(
    codcli_estado_df,
    left_index=True,
    right_on="_IDX",
    how="left",
).drop(columns=["_IDX"], errors="ignore")

base["ESTADOS_ACADEMICOS_DESDE_CODCLI"] = base["ESTADOS_ACADEMICOS_DESDE_CODCLI"].fillna("")
base["HOJAS_ESTADO_DESDE_CODCLI"] = base["HOJAS_ESTADO_DESDE_CODCLI"].fillna("")
base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI"] = base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI"].fillna(False).astype(bool)

base["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"] = (
    base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD"]
    | base["TIENE_REGLA_GOBERNADA_SIN_ACTIVIDAD_CODCLI"]
)

base["ESTADO_ACADEMICO_CONSOLIDADO"] = (
    base["ESTADOS_ACADEMICOS"].astype(str)
    + " | "
    + base["ESTADOS_ACADEMICOS_DESDE_CODCLI"].astype(str)
).str.strip(" |")

base["HOJA_ESTADO_CONSOLIDADA"] = (
    base["HOJAS_ESTADO"].astype(str)
    + " | "
    + base["HOJAS_ESTADO_DESDE_CODCLI"].astype(str)
).str.strip(" |")

base["VALORES_16_21_SIN_ACTIVIDAD_CALZAN"] = base.apply(validar_valores_sin_actividad, axis=1)

base["DICTAMEN_ANUAL_ORIGINAL"] = base["DICTAMEN_ANUAL_2025"]

def reclasificar_anual(row):
    original = str(row.get("DICTAMEN_ANUAL_ORIGINAL", ""))

    if row["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"]:
        if row["VALORES_16_21_SIN_ACTIVIDAD_CALZAN"]:
            return "OK_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD_ACADEMICA"
        return "REVISAR_EXCEPCION_GOBERNADA_VALORES_16_21_NO_CALZAN"

    return original

base["DICTAMEN_ANUAL_RECLASIFICADO"] = base.apply(reclasificar_anual, axis=1)

base["ES_BLOQUEO_REAL_POST_REGLA"] = base["DICTAMEN_ANUAL_RECLASIFICADO"].astype(str).str.contains(
    "BLOQUEO|DIFERENCIA|REVISAR",
    case=False,
    na=False,
) & ~base["DICTAMEN_ANUAL_RECLASIFICADO"].astype(str).str.startswith("OK_EXCEPCION_GOBERNADA")

base["ES_RESUELTO_POR_REGLA_GOBERNADA"] = base["DICTAMEN_ANUAL_RECLASIFICADO"].eq(
    "OK_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD_ACADEMICA"
)

base["ACCION_POST_REGLA"] = np.where(
    base["ES_RESUELTO_POR_REGLA_GOBERNADA"],
    "NO_BLOQUEAR: estado académico gobernado sin actividad académica y valores 16-21 calzan en NO/NO/0/0/0/0.",
    np.where(
        base["DICTAMEN_ANUAL_RECLASIFICADO"].eq("REVISAR_EXCEPCION_GOBERNADA_VALORES_16_21_NO_CALZAN"),
        "REVISAR: tiene estado gobernado sin actividad académica, pero las columnas 16-21 no calzan con NO/NO/0/0/0/0.",
        "MANTIENE_REVISION_ORIGINAL.",
    ),
)


# =============================================================================
# ACUMULADO
# =============================================================================

acum = aud_acum.copy()

if "RUT_NORMALIZADO" in acum.columns:
    acum = acum.merge(
        base[
            [
                "RUT_NORMALIZADO",
                "ESTADO_ACADEMICO_CONSOLIDADO",
                "HOJA_ESTADO_CONSOLIDADA",
                "TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD",
                "VALORES_16_21_SIN_ACTIVIDAD_CALZAN",
                "ES_RESUELTO_POR_REGLA_GOBERNADA",
            ]
        ].drop_duplicates("RUT_NORMALIZADO"),
        on="RUT_NORMALIZADO",
        how="left",
    )

    acum["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"] = acum["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"].fillna(False).astype(bool)
    acum["VALORES_16_21_SIN_ACTIVIDAD_CALZAN"] = acum["VALORES_16_21_SIN_ACTIVIDAD_CALZAN"].fillna(False).astype(bool)
    acum["ES_RESUELTO_POR_REGLA_GOBERNADA"] = acum["ES_RESUELTO_POR_REGLA_GOBERNADA"].fillna(False).astype(bool)

    if "DICTAMEN_ACUMULADO_2025" not in acum.columns:
        acum["DICTAMEN_ACUMULADO_2025"] = ""

    acum["DICTAMEN_ACUMULADO_RECLASIFICADO"] = np.where(
        acum["ES_RESUELTO_POR_REGLA_GOBERNADA"],
        "OK_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD_ACADEMICA",
        acum["DICTAMEN_ACUMULADO_2025"],
    )
else:
    acum["DICTAMEN_ACUMULADO_RECLASIFICADO"] = acum.get("DICTAMEN_ACUMULADO_2025", "")


# =============================================================================
# RESÚMENES
# =============================================================================

total_filas = len(base)
filas_excepcion = int(base["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"].sum())
filas_resueltas_regla = int(base["ES_RESUELTO_POR_REGLA_GOBERNADA"].sum())
filas_excepcion_no_calza = int(base["DICTAMEN_ANUAL_RECLASIFICADO"].eq("REVISAR_EXCEPCION_GOBERNADA_VALORES_16_21_NO_CALZAN").sum())
bloqueos_reales_post = int(base["ES_BLOQUEO_REAL_POST_REGLA"].sum())

if "DICTAMEN_ANUAL_ORIGINAL" in base.columns:
    bloqueos_originales = int(
        base["DICTAMEN_ANUAL_ORIGINAL"].astype(str).str.contains(
            "BLOQUEO|DIFERENCIA|REVISAR",
            case=False,
            na=False,
        ).sum()
    )
else:
    bloqueos_originales = 0

filas_ok_post = total_filas - bloqueos_reales_post

resumen_reclasificado = (
    base.groupby("DICTAMEN_ANUAL_RECLASIFICADO", dropna=False)
    .size()
    .reset_index(name="FILAS")
    .sort_values("FILAS", ascending=False)
)

resumen_estado_academico = (
    base.groupby(
        [
            "ESTADO_ACADEMICO_CONSOLIDADO",
            "TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD",
            "VALORES_16_21_SIN_ACTIVIDAD_CALZAN",
        ],
        dropna=False,
    )
    .size()
    .reset_index(name="FILAS")
    .sort_values("FILAS", ascending=False)
)

excepciones_gobernadas = base[base["TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD"]].copy()
excepciones_resueltas = base[base["ES_RESUELTO_POR_REGLA_GOBERNADA"]].copy()
excepciones_no_calzan = base[
    base["DICTAMEN_ANUAL_RECLASIFICADO"].eq("REVISAR_EXCEPCION_GOBERNADA_VALORES_16_21_NO_CALZAN")
].copy()
pendientes_reales = base[base["ES_BLOQUEO_REAL_POST_REGLA"]].copy()

dictamen_anual_post = (
    "OK_ANUAL_POST_REGLA_GOBERNADA"
    if bloqueos_reales_post == 0
    else "REVISAR_PENDIENTES_REALES_POST_REGLA_GOBERNADA"
)

dictamen_global = pd.DataFrame([
    {
        "PROCESO": PROCESO,
        "SUBPROYECTO": SUBPROYECTO,
        "ANIO_REFERENCIA": ANIO_REFERENCIA,
        "TOTAL_FILAS_5809": total_filas,
        "BLOQUEOS_O_REVISIONES_ORIGINALES": bloqueos_originales,
        "FILAS_CON_ESTADO_GOBERNADO_SIN_ACTIVIDAD": filas_excepcion,
        "FILAS_RESUELTAS_POR_REGLA_GOBERNADA": filas_resueltas_regla,
        "FILAS_EXCEPCION_GOBERNADA_NO_CALZA_16_21": filas_excepcion_no_calza,
        "BLOQUEOS_REALES_POST_REGLA": bloqueos_reales_post,
        "FILAS_OK_POST_REGLA": filas_ok_post,
        "DICTAMEN_ANUAL_POST_REGLA": dictamen_anual_post,
        "DECLARACION_CARGA": "NO_LISTO_PARA_CARGA_DEFINITIVA" if bloqueos_reales_post > 0 else "CANDIDATO_A_REVALIDACION_FINAL",
        "NOTA": "La regla ELIMINADO/VIGENTE SIN ACTIVIDAD ACADÉMICA es decisión interna gobernada; no se presenta como regla oficial del instructivo.",
    }
])

validaciones = pd.DataFrame([
    {
        "VALIDACION": "EXCEL_MAPEO_30_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(EXCEL_MAPEO_30),
    },
    {
        "VALIDACION": "AUDITORIA_CORREGIDA_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(AUDITORIA_CORREGIDA),
    },
    {
        "VALIDACION": "PROMEDIOS_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(PROMEDIOS),
    },
    {
        "VALIDACION": "CSV_5809_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(CSV_5809),
    },
    {
        "VALIDACION": "ESTADO_ACADEMICO_DETECTADO_EN_FUENTE",
        "RESULTADO": "OK" if not mapa_estado_total.empty else "REVISAR",
        "DETALLE": f"Filas mapa estado académico: {len(mapa_estado_total)}",
    },
    {
        "VALIDACION": "REGLA_GOBERNADA_APLICADA",
        "RESULTADO": "OK",
        "DETALLE": "ELIMINADO / VIGENTE SIN ACTIVIDAD ACADEMICA => excepción gobernada, no bloqueo si 16-21 calza NO/NO/0/0/0/0.",
    },
    {
        "VALIDACION": "FUENTES_ORIGINALES_MODIFICADAS",
        "RESULTADO": "OK",
        "DETALLE": "NO",
    },
])

regla_documentada = pd.DataFrame([
    {
        "NIVEL_RESPALDO": "Regla oficial",
        "FUENTE": "Instructivo_Avance Curricular SIES - 2026",
        "CONTENIDO": "Se informan presencia en 1er/2º semestre 2025, avance anual 2025 y acumulado hasta cierre académico 2025. VIGENCIA mantiene/elimina registros del proceso; todos los registros se consideran vigentes 1 aunque haya retiro posterior al 30 de abril de 2025.",
        "USO_EN_SCRIPT": "No altera VIGENCIA por estado académico institucional.",
    },
    {
        "NIVEL_RESPALDO": "Decisión interna gobernada",
        "FUENTE": "Proyecto Avance Curricular / criterio ya gobernado",
        "CONTENIDO": "Estado académico institucional ELIMINADO o VIGENTE SIN ACTIVIDAD ACADÉMICA se interpreta como caso sin actividad académica.",
        "USO_EN_SCRIPT": "No se considera bloqueo si las columnas 16-21 del 5809 son NO/NO/0/0/0/0.",
    },
    {
        "NIVEL_RESPALDO": "Implementación técnica",
        "FUENTE": "Script actual",
        "CONTENIDO": "Reclasifica dictámenes originales solo si existe estado académico gobernado y los valores 16-21 calzan con sin actividad.",
        "USO_EN_SCRIPT": "Genera salidas auditables sin modificar el 5809 original.",
    },
])

cols_export = [
    "RUT_NORMALIZADO",
    "CODCLI_LISTA",
    "CODCLI_CANTIDAD",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
    "ESTADO_ACADEMICO_CONSOLIDADO",
    "HOJA_ESTADO_CONSOLIDADA",
    "TIENE_EXCEPCION_GOBERNADA_SIN_ACTIVIDAD",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
    "VALORES_16_21_SIN_ACTIVIDAD_CALZAN",
    "DICTAMEN_ANUAL_ORIGINAL",
    "DICTAMEN_ANUAL_RECLASIFICADO",
    "ES_RESUELTO_POR_REGLA_GOBERNADA",
    "ES_BLOQUEO_REAL_POST_REGLA",
    "ACCION_POST_REGLA",
]

cols_export = [c for c in cols_export if c in base.columns]


# =============================================================================
# EXPORTAR
# =============================================================================

excel_salida = RESULTADOS / "REGLA_GOBERNADA_SIN_ACTIVIDAD_5809_2025.xlsx"

with pd.ExcelWriter(excel_salida, engine="openpyxl") as writer:
    dictamen_global.to_excel(writer, sheet_name="00_DICTAMEN_POST_REGLA", index=False)
    validaciones.to_excel(writer, sheet_name="01_VALIDACIONES", index=False)
    regla_documentada.to_excel(writer, sheet_name="02_REGLA_DOCUMENTADA", index=False)
    resumen_reclasificado.to_excel(writer, sheet_name="03_RESUMEN_RECLASIFICADO", index=False)
    resumen_estado_academico.to_excel(writer, sheet_name="04_RESUMEN_ESTADO_ACAD", index=False)
    diagnostico_estado_cols_df.to_excel(writer, sheet_name="05_DIAG_COLUMNAS_ESTADO", index=False)
    mapa_estado_total.to_excel(writer, sheet_name="06_MAPA_ESTADO_TOTAL", index=False)
    estado_por_rut.to_excel(writer, sheet_name="07_ESTADO_POR_RUT", index=False)
    estado_por_codcli.to_excel(writer, sheet_name="08_ESTADO_POR_CODCLI", index=False)
    base[cols_export].to_excel(writer, sheet_name="09_AUDITORIA_RECLASIFICADA", index=False)
    excepciones_gobernadas[cols_export].to_excel(writer, sheet_name="10_EXCEPCIONES_GOBERNADAS", index=False)
    excepciones_resueltas[cols_export].to_excel(writer, sheet_name="11_RESUELTAS_POR_REGLA", index=False)
    excepciones_no_calzan[cols_export].to_excel(writer, sheet_name="12_EXCEPCION_NO_CALZA", index=False)
    pendientes_reales[cols_export].to_excel(writer, sheet_name="13_PENDIENTES_REALES", index=False)

    if not acum.empty:
        acum.to_excel(writer, sheet_name="14_ACUMULADO_RECLASIFICADO", index=False)

    mapeo_30_p1.to_excel(writer, sheet_name="RAW_P1_MAPEO_30", index=False)
    mapeo_30_p2.to_excel(writer, sheet_name="RAW_P2_MAPEO_30", index=False)
    mapeo_30_p3.to_excel(writer, sheet_name="RAW_P3_MAPEO_30", index=False)
    mapeo_30_p4.to_excel(writer, sheet_name="RAW_P4_MAPEO_30", index=False)

    pd.DataFrame([
        {
            "TIPO": "EXCEL_MAPEO_30",
            "RUTA": str(EXCEL_MAPEO_30),
            "SHA256": sha256(EXCEL_MAPEO_30),
        },
        {
            "TIPO": "AUDITORIA_CORREGIDA",
            "RUTA": str(AUDITORIA_CORREGIDA),
            "SHA256": sha256(AUDITORIA_CORREGIDA),
        },
        {
            "TIPO": "PROMEDIOS",
            "RUTA": str(PROMEDIOS),
            "SHA256": sha256(PROMEDIOS),
        },
        {
            "TIPO": "CSV_5809",
            "RUTA": str(CSV_5809),
            "SHA256": sha256(CSV_5809),
        },
    ]).to_excel(writer, sheet_name="FUENTES", index=False)

aplicar_formato(excel_salida)
hojas_generadas = validar_xlsx(excel_salida)

informe = REPORTES / "INFORME_REGLA_GOBERNADA_SIN_ACTIVIDAD_5809_2025.md"

informe.write_text(
    "# Aplicación regla gobernada sin actividad académica 5809\n\n"
    f"Proceso: {PROCESO}  \n"
    f"Subproyecto: {SUBPROYECTO}  \n"
    f"Año referencia: {ANIO_REFERENCIA}  \n\n"
    "## Separación de respaldo\n\n"
    "- Regla oficial: el instructivo exige informar avance anual 2025 y acumulado hasta cierre 2025; la VIGENCIA del proceso no representa estado actual del estudiante.\n"
    "- Decisión interna gobernada: ELIMINADO o VIGENTE SIN ACTIVIDAD ACADÉMICA se trata como sin actividad académica.\n"
    "- Implementación técnica: si columnas 16-21 son NO/NO/0/0/0/0, el caso deja de ser bloqueo y queda como excepción gobernada.\n\n"
    "## Resultado\n\n"
    f"- Total filas 5809: {total_filas}\n"
    f"- Bloqueos/revisiones originales: {bloqueos_originales}\n"
    f"- Filas con estado gobernado sin actividad: {filas_excepcion}\n"
    f"- Filas resueltas por regla gobernada: {filas_resueltas_regla}\n"
    f"- Filas con excepción gobernada que no calzan 16-21: {filas_excepcion_no_calza}\n"
    f"- Bloqueos reales post regla: {bloqueos_reales_post}\n"
    f"- Dictamen anual post regla: {dictamen_anual_post}\n\n"
    "## Archivo\n\n"
    f"- Excel: `{excel_salida}`\n"
    f"- Carpeta repo: `{SALIDA}`\n"
    f"- Carpeta Escritorio: `{CARPETA_ESCRITORIO}`\n",
    encoding="utf-8",
)

manifest = {
    "fecha": datetime.now().isoformat(),
    "proceso": PROCESO,
    "subproyecto": SUBPROYECTO,
    "anio_referencia": ANIO_REFERENCIA,
    "excel_mapeo_30": str(EXCEL_MAPEO_30),
    "auditoria_corregida": str(AUDITORIA_CORREGIDA),
    "promedios": str(PROMEDIOS),
    "csv_5809": str(CSV_5809),
    "excel_salida": str(excel_salida),
    "informe": str(informe),
    "carpeta_repo": str(SALIDA),
    "carpeta_escritorio": str(CARPETA_ESCRITORIO),
    "total_filas_5809": total_filas,
    "bloqueos_revisiones_originales": bloqueos_originales,
    "filas_con_estado_gobernado_sin_actividad": filas_excepcion,
    "filas_resueltas_por_regla_gobernada": filas_resueltas_regla,
    "filas_excepcion_gobernada_no_calza_16_21": filas_excepcion_no_calza,
    "bloqueos_reales_post_regla": bloqueos_reales_post,
    "filas_ok_post_regla": filas_ok_post,
    "dictamen_anual_post_regla": dictamen_anual_post,
    "declaracion_carga": "NO_LISTO_PARA_CARGA_DEFINITIVA" if bloqueos_reales_post > 0 else "CANDIDATO_A_REVALIDACION_FINAL",
    "hojas_generadas": hojas_generadas,
    "fuentes_originales_modificadas": False,
}

manifest_path = MANIFEST_DIR / "manifest_regla_gobernada_sin_actividad_5809_2025.json"
manifest_path.write_text(dumps_json_safe(manifest), encoding="utf-8")

script_actual = Path(__file__).resolve()

for archivo in [excel_salida, informe, manifest_path, script_actual]:
    shutil.copy2(archivo, CARPETA_ESCRITORIO / archivo.name)

leeme = CARPETA_ESCRITORIO / "LEEME_REGLA_GOBERNADA_SIN_ACTIVIDAD_5809_2025.txt"

leeme.write_text(
    "REGLA GOBERNADA SIN ACTIVIDAD ACADEMICA 5809 - AVANCE CURRICULAR SIES 2026\n"
    f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
    f"Dictamen anual post regla: {dictamen_anual_post}\n"
    f"Bloqueos reales post regla: {bloqueos_reales_post}\n"
    f"Filas resueltas por regla gobernada: {filas_resueltas_regla}\n\n"
    "Hojas clave:\n"
    "- 00_DICTAMEN_POST_REGLA\n"
    "- 02_REGLA_DOCUMENTADA\n"
    "- 09_AUDITORIA_RECLASIFICADA\n"
    "- 10_EXCEPCIONES_GOBERNADAS\n"
    "- 11_RESUELTAS_POR_REGLA\n"
    "- 13_PENDIENTES_REALES\n\n"
    "Advertencia:\n"
    "No se modificó el 5809 original ni las fuentes. Este archivo aplica una regla interna gobernada y deja trazabilidad.\n",
    encoding="utf-8",
)

print()
print("=" * 120)
print("REGLA GOBERNADA SIN ACTIVIDAD ACADEMICA 5809 — APLICADA")
print("=" * 120)
print("Fuentes originales modificadas: NO")
print()
print("DICTAMEN POST REGLA")
print("-" * 120)
print(dictamen_global.to_string(index=False))
print()
print("RESUMEN RECLASIFICADO")
print("-" * 120)
print(resumen_reclasificado.to_string(index=False))
print()
print("VALIDACIONES")
print("-" * 120)
print(validaciones.to_string(index=False))
print()
print("ARCHIVOS")
print("-" * 120)
print(f"Excel:              {excel_salida}")
print(f"Informe:            {informe}")
print(f"Manifest:           {manifest_path}")
print(f"Script:             {script_actual}")
print(f"Carpeta repo:       {SALIDA}")
print(f"Carpeta escritorio: {CARPETA_ESCRITORIO}")
print("=" * 120)
