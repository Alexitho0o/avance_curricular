#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Auditoria especifica de los 122 casos SIN_REGISTROS_2025_PARA_CODCLI_LISTA.

No modifica fuentes originales, no corrige datos, no genera SIES_READY,
no recalcula 20-21 y no mezcla hitos 54 ni 55. Solo diagnostica causas y
propone rutas de resolucion.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Auditoria sin registros 2025 CODCLI_LISTA 5809"
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
NO_CORRECCION = "NO_APLICAR_CORRECCION_AUTOMATICA"
ESTADO_VALIDADOR_REQUERIDO = "APTO_PARA_CONTROL_PREVIO"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_56 = REPO / "avance_curricular_2026" / "56_auditoria_sin_registros_2025_codcli_lista_5809"

GOBERNANZA_51 = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)
VALIDADOR_52_SCRIPT = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "validador_gobernanza_precalculo_5809.py"
)
VALIDADOR_52_XLSX = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx"
)
MATRIZ_53 = (
    REPO
    / "avance_curricular_2026/53_matriz_resolucion_bloqueos_16_19_gobernada_5809/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_20260708_170521/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_5809.xlsx"
)
RECALCULO_50 = (
    REPO
    / "avance_curricular_2026/50_recalculo_16_19_diccionario_estado_promedios/"
    / "RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS_20260708_143748/"
    / "02_RESULTADOS/RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS.xlsx"
)
PROMEDIOS = (
    REPO
    / "avance_curricular_2026/25_actualizacion_fuente_promedios/"
    / "CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/00_FUENTE_CONGELADA/"
    / "PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
)
MAPEO_IDENTIDAD = (
    REPO
    / "avance_curricular_2026/05_cierre_integral/CIERRE_AVANCE_CURRICULAR_20260703_125157/"
    / "04_CONCILIACION_MATRICULA/PREPARACION_MATRICULA_20260703_141254/02_IDENTIFICADORES/"
    / "MAPEO_IDENTIDAD_CODCLI.xlsx"
)

OUTPUT_FILES = {
    "excel": "AUDITORIA_SIN_REGISTROS_2025_CODCLI_LISTA_5809.xlsx",
    "informe": "INFORME_AUDITORIA_SIN_REGISTROS_2025_CODCLI_LISTA_5809.md",
    "manifest": "manifest_auditoria_sin_registros_2025_codcli_lista_5809.json",
    "script": "auditoria_sin_registros_2025_codcli_lista_5809.py",
}

CAUSES = [
    "CODCLI_EXISTE_SOLO_EN_2026",
    "CODCLI_EXISTE_SOLO_OTRO_ANO",
    "CODCLI_NO_EXISTE_EN_PROMEDIOS",
    "RUT_TIENE_OTRO_CODCLI_CON_2025",
    "ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025",
    "POSIBLE_MAPEO_CODCLI_LISTA_INCOMPLETO",
    "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS",
    "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA",
    "REQUIERE_REVISION_FUNCIONAL",
]

INACTIVE_STATUS = {"ELIMINADO", "SUSPENDIDO", "TITULADO", "RETIRADO", "ANULADO"}

SHEETS = OrderedDict(
    [
        ("00_DICTAMEN", "00_DICTAMEN"),
        ("01_VALIDACION_GOBERNANZA", "01_VALIDACION_GOBERNANZA"),
        ("02_RESUMEN_CAUSAS", "02_RESUMEN_CAUSAS"),
        ("03_122_CASOS", "03_122_CASOS"),
        ("04_BUSQUEDA_CODCLI_OTROS_ANOS", "04_BUSQUEDA_CODCLI_OTROS_ANOS"),
        ("05_BUSQUEDA_RUT_PROMEDIOS", "05_BUSQUEDA_RUT_PROMEDIOS"),
        ("06_ESTADO_ACADEMICO_OBSERVADO", "06_ESTADO_ACADEMICO_OBSERVADO"),
        ("07_MAPEO_CODCLI_LISTA", "07_MAPEO_CODCLI_LISTA"),
        ("08_CAUSAS_PROPUESTAS", "08_CAUSAS_PROPUESTAS"),
        ("09_NO_RESOLVER_AUTOMATICO", "09_NO_RESOLVER_AUTOMATICO"),
        ("10_FUENTES", "10_FUENTES"),
    ]
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def norm(value: Any) -> str:
    text = clean(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    out = []
    for ch in text.upper():
        out.append(ch if ch.isalnum() else "_")
    return "_".join("".join(out).split("_"))


def compact_number(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return text


def display(value: Any, max_len: int = 160) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")


def safe_read_excel(path: Path, sheet: str, columns: Iterable[str] | None = None) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")
    if columns is None:
        return df
    existing = [col for col in columns if col in df.columns]
    return df[existing].copy()


def split_codcli_lista(value: Any) -> List[str]:
    text = clean(value)
    if not text:
        return []
    text = text.replace("|", ";")
    return [part.strip() for part in text.split(";") if part.strip()]


def join_unique(values: Iterable[Any], sep: str = " | ") -> str:
    cleaned = sorted(set(clean(v) for v in values if clean(v)))
    return sep.join(cleaned)


def to_int(value: Any) -> int:
    text = clean(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def is_si(value: Any) -> bool:
    return norm(value) == "SI"


def normalize_rut_body(value: Any) -> str:
    return "".join(ch for ch in clean(value).upper() if ch.isalnum())


def normalize_rut_full(body: Any, dv: Any = "") -> str:
    body_text = normalize_rut_body(body)
    dv_text = normalize_rut_body(dv)
    if "-" in clean(body) and not dv_text:
        return clean(body).upper()
    if body_text and dv_text:
        return f"{body_text}-{dv_text}"
    return body_text


def require_paths() -> None:
    paths = [
        GOBERNANZA_51,
        VALIDADOR_52_SCRIPT,
        VALIDADOR_52_XLSX,
        MATRIZ_53,
        RECALCULO_50,
        PROMEDIOS,
        MAPEO_IDENTIDAD,
    ]
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        raise FileNotFoundError("Faltan entradas obligatorias:\n" + "\n".join(missing))


def run_validator_52(output_dir: Path, timestamp: str) -> Dict[str, Any]:
    plan_path = output_dir / "plan_uso_validador_h52_auditoria_sin_registros_2025.json"
    plan = {
        "descripcion": (
            "Auditoria diagnostica de 122 casos SIN_REGISTROS_2025_PARA_CODCLI_LISTA. "
            "No corrige, no recalcula 20-21, no genera carga y no mezcla hitos 54/55."
        ),
        "columnas_usadas": ["CODIGO_UNICO", "PLAN_ESTUDIOS", "CODCLI_LISTA", "CODCLI", "ANO"],
        "llaves_usadas": ["CODCLI_LISTA", "CODIGO_UNICO", "PLAN_ESTUDIOS"],
        "diccionarios_usados": [],
        "transformaciones": [
            "Diagnostico documental de ausencia de registros ANO=2025 para CODCLI_LISTA; sin correccion de datos"
        ],
        "campos_calculo": [],
        "anio_calculo": 2025,
        "genera_sies_ready": False,
        "modifica_fuentes_originales": False,
        "usa_todos_codcli_rut": False,
        "filtra_por_programa": True,
        "usa_inferencias_como_reglas_oficiales": False,
        "usa_estados_AEIR": False,
        "usa_codcli_lista": True,
        "mezcla_anual_acumulado": False,
        "modo": "AUDITORIA_DIAGNOSTICA_NO_CORRECTIVA_SIN_REGISTROS_2025",
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")

    validator_timestamp = f"{timestamp}_H56"
    cmd = [
        sys.executable,
        str(VALIDADOR_52_SCRIPT),
        "--timestamp",
        validator_timestamp,
        "--plan-json",
        str(plan_path),
    ]
    completed = subprocess.run(cmd, check=False, capture_output=True, text=True)
    validator_dir = (
        REPO
        / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809"
        / f"VALIDADOR_GOBERNANZA_PRECALCULO_5809_{validator_timestamp}"
    )
    manifest_path = validator_dir / "manifest_validacion_gobernanza_precalculo_5809.json"
    manifest: Dict[str, Any] = {}
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    result = {
        "plan_path": str(plan_path),
        "validator_timestamp": validator_timestamp,
        "validator_dir": str(validator_dir),
        "validator_manifest": str(manifest_path),
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "estado_validador": manifest.get("estado_validador", "SIN_MANIFEST"),
        "manifest": manifest,
    }
    if completed.returncode != 0:
        raise RuntimeError(
            "Validador hito 52 fallo.\nSTDOUT:\n"
            + completed.stdout
            + "\nSTDERR:\n"
            + completed.stderr
        )
    if result["estado_validador"] != ESTADO_VALIDADOR_REQUERIDO:
        raise RuntimeError(
            f"Validador hito 52 no apto: {result['estado_validador']}. "
            f"Manifest: {manifest_path}"
        )
    return result


def prepare_promedios() -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    hoja1_cols = [
        "CODCLI",
        "RUT",
        "DIG",
        "ANO",
        "PERIODO",
        "ESTADO",
        "DESCRIPCION_ESTADO",
        "CODRAMO",
        "ESTADO_ACADEMICO",
        "PLAN_DE_ESTUDIO",
    ]
    alumnos_cols = [
        "CODCLI",
        "CODCARPR",
        "NOMBRE_L",
        "ANOMATRICULA",
        "PERIODOMATRICULA",
        "ANOINGRESO",
        "PERIODOINGRESO",
        "RUT",
        "ESTADOACADEMICO",
        "SITUACION",
        "MATRICULA",
        "NIVEL",
    ]
    hoja1 = safe_read_excel(PROMEDIOS, "Hoja1", hoja1_cols)
    datos = safe_read_excel(PROMEDIOS, "DatosAlumnos", alumnos_cols)
    matricula = safe_read_excel(PROMEDIOS, "Matricula_2025", alumnos_cols)

    for df in [hoja1, datos, matricula]:
        if "CODCLI" in df.columns:
            df["CODCLI_N"] = df["CODCLI"].map(clean)
        if "RUT" in df.columns:
            df["RUT_BODY_N"] = df["RUT"].map(normalize_rut_body)

    if "ANO" in hoja1.columns:
        hoja1["ANO_N"] = hoja1["ANO"].map(compact_number)
    else:
        hoja1["ANO_N"] = ""
    hoja1["RUT_FULL_N"] = [
        normalize_rut_full(row.get("RUT", ""), row.get("DIG", "")) for _, row in hoja1.iterrows()
    ]
    for df in [datos, matricula]:
        if "ANOMATRICULA" in df.columns:
            df["ANOMATRICULA_N"] = df["ANOMATRICULA"].map(compact_number)
    return hoja1, datos, matricula


def read_mapeo_for_cases(cases: pd.DataFrame) -> pd.DataFrame:
    mapeo = read_sheet(MAPEO_IDENTIDAD, "MAPEO")
    if "ID_FILA_5809" not in mapeo.columns:
        return pd.DataFrame()
    case_ids = {clean(v) for v in cases["FILA_5809"].tolist()}
    out = mapeo[mapeo["ID_FILA_5809"].map(clean).isin(case_ids)].copy()
    return out.fillna("")


def status_for_codcli(source: str, codcli: str, df: pd.DataFrame) -> Dict[str, Any]:
    if df.empty or "CODCLI_N" not in df.columns:
        return {
            "Fuente estado": source,
            "CODCLI": codcli,
            "Filas fuente": 0,
            "Estados academicos": "",
            "Situaciones": "",
            "Matriculas": "",
            "Anos matricula": "",
        }
    sub = df[df["CODCLI_N"] == codcli].copy()
    estado_values: List[Any] = []
    for col in ["ESTADO_ACADEMICO", "ESTADOACADEMICO"]:
        if col in sub.columns:
            estado_values.extend(sub[col].tolist())
    return {
        "Fuente estado": source,
        "CODCLI": codcli,
        "Filas fuente": len(sub),
        "Estados academicos": join_unique(estado_values),
        "Situaciones": join_unique(sub["SITUACION"]) if "SITUACION" in sub.columns else "",
        "Matriculas": join_unique(sub["MATRICULA"]) if "MATRICULA" in sub.columns else "",
        "Anos matricula": join_unique(sub["ANOMATRICULA_N"]) if "ANOMATRICULA_N" in sub.columns else "",
    }


def any_inactive(statuses: Iterable[Any]) -> bool:
    normalized = {norm(value) for value in statuses if clean(value)}
    return bool(normalized & INACTIVE_STATUS)


def case_has_5809_activity(row: pd.Series) -> bool:
    return any(
        [
            is_si(row.get("CURSO_1ER_SEM_5809", "")),
            is_si(row.get("CURSO_2DO_SEM_5809", "")),
            to_int(row.get("UNIDADES_CURSADAS_5809", "")) > 0,
            to_int(row.get("UNIDADES_APROBADAS_5809", "")) > 0,
        ]
    )


def classify_case(
    row: pd.Series,
    codcli_list: List[str],
    hoja1: pd.DataFrame,
    datos: pd.DataFrame,
    matricula: pd.DataFrame,
    mapeo_row: pd.DataFrame,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]], Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    fila = clean(row.get("FILA_5809", ""))
    rut_normalizado = clean(row.get("RUT_NORMALIZADO", ""))
    codcli_set = set(codcli_list)

    codcli_sub = hoja1[hoja1["CODCLI_N"].isin(codcli_set)].copy() if codcli_set else hoja1.iloc[0:0].copy()
    years = sorted(set(v for v in codcli_sub["ANO_N"].tolist() if clean(v)))
    rows_2025 = codcli_sub[codcli_sub["ANO_N"] == "2025"]
    rows_2026 = codcli_sub[codcli_sub["ANO_N"] == "2026"]
    rows_other_year = codcli_sub[(codcli_sub["ANO_N"] != "") & (~codcli_sub["ANO_N"].isin(["2025", "2026"]))]
    rows_blank_year = codcli_sub[codcli_sub["ANO_N"] == ""]

    rut_body = rut_normalizado.split("-")[0] if "-" in rut_normalizado else normalize_rut_body(rut_normalizado)
    rut_sub = hoja1[hoja1["RUT_BODY_N"] == rut_body].copy() if rut_body else hoja1.iloc[0:0].copy()
    rut_2025 = rut_sub[rut_sub["ANO_N"] == "2025"].copy()
    rut_2025_other_codcli = sorted(set(rut_2025["CODCLI_N"].tolist()) - codcli_set)

    estado_rows: List[Dict[str, Any]] = []
    estados_observados: List[str] = []
    for codcli in codcli_list:
        h1_status = status_for_codcli("PROMEDIOS_Hoja1", codcli, codcli_sub)
        da_status = status_for_codcli("PROMEDIOS_DatosAlumnos", codcli, datos)
        m25_status = status_for_codcli("PROMEDIOS_Matricula_2025", codcli, matricula)
        for status in [h1_status, da_status, m25_status]:
            status["FILA_5809"] = fila
            status["RUT_NORMALIZADO"] = rut_normalizado
            estado_rows.append(status)
            estados_observados.extend(split_codcli_lista(status.get("Estados academicos", "").replace(" | ", ";")))

    inactive = any_inactive(estados_observados)
    codcli_count = to_int(row.get("CODCLI_LISTA_CANTIDAD", len(codcli_list))) or len(codcli_list)
    mapeo_clas = join_unique(mapeo_row.get("CLASIFICACION_IDENTIDAD", [])) if not mapeo_row.empty else ""
    mapeo_incomplete = codcli_count > 1 or (mapeo_clas and norm(mapeo_clas) != "IDENTIDAD_EXACTA")
    has_activity = case_has_5809_activity(row)

    flags = {
        "CODCLI_EXISTE_SOLO_EN_2026": bool(not rows_2026.empty and len(years) == 1 and years[0] == "2026"),
        "CODCLI_EXISTE_SOLO_OTRO_ANO": bool(not codcli_sub.empty and bool(years) and "2025" not in years and years != ["2026"]),
        "CODCLI_NO_EXISTE_EN_PROMEDIOS": bool(codcli_sub.empty),
        "RUT_TIENE_OTRO_CODCLI_CON_2025": bool(rut_2025_other_codcli),
        "ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025": bool(inactive),
        "POSIBLE_MAPEO_CODCLI_LISTA_INCOMPLETO": bool(mapeo_incomplete),
        "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS": bool(has_activity and rows_2025.empty),
        "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA": bool(rows_2025.empty),
        "REQUIERE_REVISION_FUNCIONAL": True,
    }

    if flags["RUT_TIENE_OTRO_CODCLI_CON_2025"]:
        principal = "RUT_TIENE_OTRO_CODCLI_CON_2025"
    elif flags["CODCLI_NO_EXISTE_EN_PROMEDIOS"]:
        principal = "CODCLI_NO_EXISTE_EN_PROMEDIOS"
    elif flags["CODCLI_EXISTE_SOLO_EN_2026"]:
        principal = "CODCLI_EXISTE_SOLO_EN_2026"
    elif flags["CODCLI_EXISTE_SOLO_OTRO_ANO"]:
        principal = "CODCLI_EXISTE_SOLO_OTRO_ANO"
    elif flags["ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025"]:
        principal = "ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025"
    elif flags["REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA"]:
        principal = "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA"
    elif flags["POSIBLE_MAPEO_CODCLI_LISTA_INCOMPLETO"]:
        principal = "POSIBLE_MAPEO_CODCLI_LISTA_INCOMPLETO"
    else:
        principal = "REQUIERE_REVISION_FUNCIONAL"

    detected = [cause for cause in CAUSES if flags[cause]]
    case_result = {
        "FILA_5809": fila,
        "RUT_NORMALIZADO": rut_normalizado,
        "CODIGO_UNICO": clean(row.get("CODIGO_UNICO", "")),
        "PLAN_ESTUDIOS": clean(row.get("PLAN_ESTUDIOS", "")),
        "CODCLI_LISTA_NORM": clean(row.get("CODCLI_LISTA_NORM", "")),
        "CODCLI_LISTA_CANTIDAD": codcli_count,
        "DICTAMEN_HITO_53": clean(row.get("DICTAMEN_DICCIONARIO_ESTADO", "")),
        "CAUSA_PRINCIPAL_PROPUESTA": principal,
        "CAUSAS_DETECTADAS": " | ".join(detected),
        "ANOS_CODCLI_LISTA_EN_PROMEDIOS": join_unique(years),
        "FILAS_PROMEDIOS_CODCLI_LISTA": len(codcli_sub),
        "FILAS_2025_CODCLI_LISTA": len(rows_2025),
        "FILAS_2026_CODCLI_LISTA": len(rows_2026),
        "FILAS_OTROS_ANOS_CODCLI_LISTA": len(rows_other_year),
        "FILAS_ANO_BLANCO_CODCLI_LISTA": len(rows_blank_year),
        "RUT_TIENE_REGISTROS_PROMEDIOS": "SI" if not rut_sub.empty else "NO",
        "ANOS_RUT_EN_PROMEDIOS": join_unique(rut_sub["ANO_N"]) if not rut_sub.empty else "",
        "CODCLI_RUT_2025_DISTINTO_LISTA": " | ".join(rut_2025_other_codcli),
        "ESTADOS_ACADEMICOS_OBSERVADOS": join_unique(estados_observados),
        "ESTADO_ACADEMICO_INACTIVO_OBSERVADO": "SI" if inactive else "NO",
        "CLASIFICACION_MAPEO": mapeo_clas,
        "MAPEO_CODCLI_LISTA_MULTIPLE_O_NO_EXACTO": "SI" if mapeo_incomplete else "NO",
        "5809_DECLARA_ACTIVIDAD_ANUAL": "SI" if has_activity else "NO",
        "ACCION_PROPUESTA": (
            "No corregir automaticamente. Solicitar evidencia academica 2025 o decision funcional documentada "
            "para mantener ausencia de registros 2025."
        ),
        "NO_APLICAR_CORRECCION": NO_CORRECCION,
        "DECLARACION_CARGA": DECLARACION_CARGA,
    }

    codcli_year_rows: List[Dict[str, Any]] = []
    for codcli in codcli_list:
        sub = hoja1[hoja1["CODCLI_N"] == codcli].copy()
        codcli_year_rows.append(
            {
                "FILA_5809": fila,
                "RUT_NORMALIZADO": rut_normalizado,
                "CODCLI": codcli,
                "FILAS_HOJA1": len(sub),
                "ANOS_OBSERVADOS": join_unique(sub["ANO_N"]) if not sub.empty else "",
                "FILAS_2025": int((sub["ANO_N"] == "2025").sum()) if not sub.empty else 0,
                "FILAS_2026": int((sub["ANO_N"] == "2026").sum()) if not sub.empty else 0,
                "FILAS_OTROS_ANOS": int(((sub["ANO_N"] != "") & (~sub["ANO_N"].isin(["2025", "2026"]))).sum())
                if not sub.empty
                else 0,
                "FILAS_ANO_BLANCO": int((sub["ANO_N"] == "").sum()) if not sub.empty else 0,
                "PERIODOS_OBSERVADOS": join_unique(sub["PERIODO"]) if "PERIODO" in sub.columns and not sub.empty else "",
                "CODRAMOS_DISTINTOS": sub["CODRAMO"].nunique() if "CODRAMO" in sub.columns and not sub.empty else 0,
                "ESTADOS_ASIGNATURA": join_unique(sub["ESTADO"]) if "ESTADO" in sub.columns and not sub.empty else "",
                "DESCRIPCIONES_ESTADO": join_unique(sub["DESCRIPCION_ESTADO"])
                if "DESCRIPCION_ESTADO" in sub.columns and not sub.empty
                else "",
                "ESTADOS_ACADEMICOS_HOJA1": join_unique(sub["ESTADO_ACADEMICO"])
                if "ESTADO_ACADEMICO" in sub.columns and not sub.empty
                else "",
                "DICTAMEN": "SIN_REGISTROS_2025_PARA_CODCLI_LISTA",
            }
        )

    rut_result = {
        "FILA_5809": fila,
        "RUT_NORMALIZADO": rut_normalizado,
        "FILAS_PROMEDIOS_RUT": len(rut_sub),
        "ANOS_RUT_EN_PROMEDIOS": join_unique(rut_sub["ANO_N"]) if not rut_sub.empty else "",
        "CODCLI_RUT_PROMEDIOS": join_unique(rut_sub["CODCLI_N"]) if not rut_sub.empty else "",
        "CODCLI_RUT_CON_2025": join_unique(rut_2025["CODCLI_N"]) if not rut_2025.empty else "",
        "CODCLI_RUT_2025_DISTINTO_LISTA": " | ".join(rut_2025_other_codcli),
        "RUT_TIENE_OTRO_CODCLI_CON_2025": "SI" if rut_2025_other_codcli else "NO",
        "USO": "Evidencia diagnostica; RUT esta no gobernado para calculo en hito 51.",
    }

    no_auto = [
        {
            "FILA_5809": fila,
            "CODCLI_LISTA_NORM": clean(row.get("CODCLI_LISTA_NORM", "")),
            "CAUSA_PRINCIPAL_PROPUESTA": principal,
            "MOTIVO_NO_AUTOMATICO": (
                "No hay registros ANO=2025 para CODCLI_LISTA en PROMEDIOS; cualquier ajuste requiere "
                "fuente academica complementaria o decision funcional."
            ),
            "NO_APLICAR_CORRECCION": NO_CORRECCION,
            "DECLARACION_CARGA": DECLARACION_CARGA,
        }
    ]
    return case_result, codcli_year_rows, rut_result, estado_rows, no_auto


def build_case_tables(
    cases: pd.DataFrame,
    hoja1: pd.DataFrame,
    datos: pd.DataFrame,
    matricula: pd.DataFrame,
    mapeo: pd.DataFrame,
) -> Dict[str, pd.DataFrame]:
    case_rows: List[Dict[str, Any]] = []
    codcli_year_rows: List[Dict[str, Any]] = []
    rut_rows: List[Dict[str, Any]] = []
    estado_rows: List[Dict[str, Any]] = []
    no_auto_rows: List[Dict[str, Any]] = []

    mapeo_by_fila = {clean(row["ID_FILA_5809"]): pd.DataFrame([row]) for _, row in mapeo.iterrows()} if not mapeo.empty else {}

    for _, row in cases.iterrows():
        fila = clean(row.get("FILA_5809", ""))
        codcli_list = split_codcli_lista(row.get("CODCLI_LISTA_NORM", ""))
        mapeo_row = mapeo_by_fila.get(fila, pd.DataFrame())
        case_result, codcli_year, rut_result, estado, no_auto = classify_case(
            row, codcli_list, hoja1, datos, matricula, mapeo_row
        )
        case_rows.append(case_result)
        codcli_year_rows.extend(codcli_year)
        rut_rows.append(rut_result)
        estado_rows.extend(estado)
        no_auto_rows.extend(no_auto)

    return {
        "casos": pd.DataFrame(case_rows),
        "codcli_otros_anos": pd.DataFrame(codcli_year_rows),
        "rut_promedios": pd.DataFrame(rut_rows),
        "estado_observado": pd.DataFrame(estado_rows),
        "no_auto": pd.DataFrame(no_auto_rows),
    }


def build_resumen_causas(cases_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for cause in CAUSES:
        principal = int((cases_df["CAUSA_PRINCIPAL_PROPUESTA"] == cause).sum())
        detected = int(cases_df["CAUSAS_DETECTADAS"].map(lambda text: cause in split_codcli_lista(clean(text))).sum())
        rows.append(
            {
                "CAUSA": cause,
                "CASOS_CAUSA_PRINCIPAL": principal,
                "CASOS_CAUSA_DETECTADA": detected,
                "RESOLVER_AUTOMATICAMENTE": "NO",
                "RUTA_PROPUESTA": route_for_cause(cause),
                "DECLARACION_CARGA": DECLARACION_CARGA,
            }
        )
    return pd.DataFrame(rows)


def route_for_cause(cause: str) -> str:
    routes = {
        "CODCLI_EXISTE_SOLO_EN_2026": "Confirmar si el CODCLI corresponde a ingreso/actividad 2026 y excluir de calculo anual 2025 sin corregir aun.",
        "CODCLI_EXISTE_SOLO_OTRO_ANO": "Revisar cobertura historica y confirmar ausencia de detalle 2025 para el programa.",
        "CODCLI_NO_EXISTE_EN_PROMEDIOS": "Solicitar fuente academica complementaria o confirmar mapeo incompleto.",
        "RUT_TIENE_OTRO_CODCLI_CON_2025": "Revisar identidad/programa antes de considerar cualquier remapeo; no sumar todos los CODCLI del RUT.",
        "ESTADO_ACADEMICO_EXPLICA_AUSENCIA_2025": "Validar funcionalmente si estado academico inactivo explica ausencia de actividad 2025.",
        "POSIBLE_MAPEO_CODCLI_LISTA_INCOMPLETO": "Revisar MAPEO_IDENTIDAD_CODCLI y programa antes de usar otro CODCLI.",
        "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS": "Contrastar 5809 con evidencia academica 2025 externa antes de ajustar.",
        "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA": "Solicitar detalle academico 2025 complementario o acta de decision funcional.",
        "REQUIERE_REVISION_FUNCIONAL": "Resolver con responsable funcional antes de cualquier correccion.",
    }
    return routes[cause]


def build_causas_propuestas(resumen: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, row in resumen.iterrows():
        rows.append(
            {
                "CAUSA": row["CAUSA"],
                "CASOS_CAUSA_PRINCIPAL": row["CASOS_CAUSA_PRINCIPAL"],
                "CASOS_CAUSA_DETECTADA": row["CASOS_CAUSA_DETECTADA"],
                "DECISION_PROPUESTA": route_for_cause(row["CAUSA"]),
                "PUEDE_RESOLVERSE_AUTOMATICAMENTE_AHORA": "NO",
                "REQUIERE_REVISION_FUNCIONAL": "SI",
                "NO_APLICAR_CORRECCION": NO_CORRECCION,
            }
        )
    return pd.DataFrame(rows)


def build_validation_df(validator: Dict[str, Any]) -> pd.DataFrame:
    manifest = validator.get("manifest", {})
    return pd.DataFrame(
        [
            {
                "Control": "Gobernanza integral hito 51",
                "Estado": "USADA",
                "Evidencia": str(GOBERNANZA_51),
                "Hash SHA256": sha256_file(GOBERNANZA_51),
                "Observacion": "Base normativa/documental obligatoria.",
            },
            {
                "Control": "Validador hito 52",
                "Estado": validator["estado_validador"],
                "Evidencia": validator["validator_manifest"],
                "Hash SHA256": sha256_file(Path(validator["validator_manifest"]))
                if Path(validator["validator_manifest"]).exists()
                else "",
                "Observacion": "Compuerta ejecutada antes de generar la auditoria hito 56.",
            },
            {
                "Control": "Matriz hito 53",
                "Estado": "USADA",
                "Evidencia": str(MATRIZ_53),
                "Hash SHA256": sha256_file(MATRIZ_53),
                "Observacion": "Se usa solo hoja 05_SIN_REGISTROS_2025; no se mezclan hitos 54/55.",
            },
            {
                "Control": "Recalculo paso 50",
                "Estado": "USADO_COMO_EVIDENCIA",
                "Evidencia": str(RECALCULO_50),
                "Hash SHA256": sha256_file(RECALCULO_50),
                "Observacion": "Evidencia tecnica no normativa; no se recalcula 20-21.",
            },
            {
                "Control": "PROMEDIOS actualizado",
                "Estado": "USADO_COMO_EVIDENCIA",
                "Evidencia": str(PROMEDIOS),
                "Hash SHA256": sha256_file(PROMEDIOS),
                "Observacion": "Busqueda de CODCLI_LISTA por anos y estados academicos observados.",
            },
            {
                "Control": "MAPEO_IDENTIDAD_CODCLI",
                "Estado": "USADO",
                "Evidencia": str(MAPEO_IDENTIDAD),
                "Hash SHA256": sha256_file(MAPEO_IDENTIDAD),
                "Observacion": "CODCLI_LISTA es columna correcta; N_CODCLI no se usa como CODCLI academico.",
            },
            {
                "Control": "Plan hito 52",
                "Estado": manifest.get("estado_validador", validator["estado_validador"]),
                "Evidencia": validator["plan_path"],
                "Hash SHA256": sha256_file(Path(validator["plan_path"])),
                "Observacion": "Plan diagnostico no correctivo, sin SIES_READY y sin modificar fuentes.",
            },
        ]
    )


def build_fuentes_df(output_dir: Path, output_paths: Dict[str, Path], validator: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    input_sources = [
        ("INPUT", "Gobernanza integral hito 51", GOBERNANZA_51),
        ("INPUT", "Validador hito 52 xlsx base", VALIDADOR_52_XLSX),
        ("INPUT", "Validador hito 52 script", VALIDADOR_52_SCRIPT),
        ("INPUT", "Matriz hito 53", MATRIZ_53),
        ("INPUT", "Recalculo paso 50", RECALCULO_50),
        ("INPUT", "PROMEDIOS actualizado", PROMEDIOS),
        ("INPUT", "MAPEO_IDENTIDAD_CODCLI", MAPEO_IDENTIDAD),
        ("INPUT_DERIVADO", "Manifest validador hito 52 ejecucion hito 56", Path(validator["validator_manifest"])),
    ]
    for tipo, nombre, path in input_sources:
        rows.append(
            {
                "Tipo": tipo,
                "Nombre": nombre,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() else "",
                "Uso": "Entrada o compuerta obligatoria; no modificada.",
            }
        )
    for key, path in output_paths.items():
        output_hash = ""
        if path.exists() and key != "manifest":
            output_hash = sha256_file(path)
        rows.append(
            {
                "Tipo": "OUTPUT",
                "Nombre": key,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": output_hash,
                "Uso": "Producto hito 56 no apto para carga. Manifest sin hash autorreferencial.",
            }
        )
    rows.append(
        {
            "Tipo": "OUTPUT_DIR",
            "Nombre": "Carpeta hito 56",
            "Ruta": str(output_dir),
            "Existe": "SI" if output_dir.exists() else "NO",
            "Hash SHA256": "",
            "Uso": "Contenedor de auditoria especifica.",
        }
    )
    return pd.DataFrame(rows)


def write_excel(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for logical_name, df in sheets.items():
            sheet_name = SHEETS.get(logical_name, logical_name)[:31]
            df.to_excel(writer, sheet_name=sheet_name, index=False)

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            values = [display(cell.value, 90) for cell in column_cells[:200]]
            width = min(max([len(v) for v in values] + [10]) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    wb.save(path)


def build_markdown(
    output_dir: Path,
    dictamen_df: pd.DataFrame,
    validacion_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    casos_df: pd.DataFrame,
    manifest: Dict[str, Any],
) -> str:
    principal = (
        resumen_df[resumen_df["CASOS_CAUSA_PRINCIPAL"] > 0]
        [["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RUTA_PROPUESTA"]]
        .copy()
    )
    lines = [
        "# Informe Auditoria SIN_REGISTROS_2025_PARA_CODCLI_LISTA 5809",
        "",
        f"- Proceso: {PROCESO}",
        f"- Subproyecto: {SUBPROYECTO}",
        f"- Anio proceso: {ANIO_PROCESO}",
        f"- Anio referencia datos: {ANIO_REFERENCIA_DATOS}",
        f"- Declaracion carga: {DECLARACION_CARGA}",
        "- Fuentes originales modificadas: NO",
        "- Correcciones aplicadas: NO",
        "- SIES_READY generado: NO",
        "- Recalculo 20-21: NO",
        "- Mezcla con hitos 54/55: NO",
        f"- Carpeta: {output_dir}",
        "",
        "## Dictamen",
        "",
        markdown_table(dictamen_df),
        "",
        "## Validacion de gobernanza",
        "",
        markdown_table(validacion_df[["Control", "Estado", "Observacion"]]),
        "",
        "## Resumen de causas",
        "",
        markdown_table(resumen_df[["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RESOLVER_AUTOMATICAMENTE"]]),
        "",
        "## Causas principales con ruta",
        "",
        markdown_table(principal),
        "",
        "## Hallazgos",
        "",
        "- Los 122 casos provienen exclusivamente de la hoja 05_SIN_REGISTROS_2025 del hito 53.",
        "- No se usan N_CODCLI ni RUT como llave academica para corregir; RUT se muestra solo como evidencia diagnostica.",
        "- No hay registros ANO=2025 asociados a CODCLI_LISTA en PROMEDIOS para estos casos.",
        "- Toda ruta propuesta queda no automatica y requiere fuente academica complementaria o decision funcional.",
        "",
        "## Muestra de casos",
        "",
        markdown_table(
            casos_df[
                [
                    "FILA_5809",
                    "CODCLI_LISTA_NORM",
                    "CAUSA_PRINCIPAL_PROPUESTA",
                    "ANOS_CODCLI_LISTA_EN_PROMEDIOS",
                    "ESTADOS_ACADEMICOS_OBSERVADOS",
                    "5809_DECLARA_ACTIVIDAD_ANUAL",
                ]
            ],
            max_rows=20,
        ),
        "",
        "## Archivos",
        "",
        f"- Excel: {manifest.get('archivos', {}).get('excel', '')}",
        f"- Informe: {manifest.get('archivos', {}).get('informe', '')}",
        f"- Manifest: {manifest.get('archivos', {}).get('manifest', '')}",
        f"- Script: {manifest.get('archivos', {}).get('script', '')}",
        "",
    ]
    return "\n".join(lines)


def build_manifest(
    timestamp: str,
    output_dir: Path,
    output_paths: Dict[str, Path],
    validator: Dict[str, Any],
    dictamen_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    sheets: Dict[str, pd.DataFrame],
) -> Dict[str, Any]:
    return {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": timestamp,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "correcciones_aplicadas": "NO",
        "sies_ready_generado": "NO",
        "recalculo_20_21": "NO",
        "mezcla_hitos_54_55": "NO",
        "estado_validador_hito_52": validator["estado_validador"],
        "carpeta_salida": str(output_dir),
        "archivos": {key: str(path) for key, path in output_paths.items()},
        "entradas_obligatorias": {
            "gobernanza_51": str(GOBERNANZA_51),
            "validador_52_script": str(VALIDADOR_52_SCRIPT),
            "validador_52_xlsx": str(VALIDADOR_52_XLSX),
            "matriz_53": str(MATRIZ_53),
            "recalculo_50": str(RECALCULO_50),
            "promedios_actualizado": str(PROMEDIOS),
            "mapeo_identidad_codcli": str(MAPEO_IDENTIDAD),
        },
        "validador_hito_52_ejecucion": {
            "plan": validator["plan_path"],
            "manifest": validator["validator_manifest"],
            "carpeta": validator["validator_dir"],
            "returncode": validator["returncode"],
        },
        "conteos": {
            "casos_auditados": int(dictamen_df.loc[0, "Total casos auditados"]),
            "casos_hoja_05_hito_53": int(dictamen_df.loc[0, "Casos fuente hito 53"]),
            "causas_principales": resumen_df.set_index("CAUSA")["CASOS_CAUSA_PRINCIPAL"].astype(int).to_dict(),
            "causas_detectadas": resumen_df.set_index("CAUSA")["CASOS_CAUSA_DETECTADA"].astype(int).to_dict(),
            "hojas_excel": {name: len(df) for name, df in sheets.items()},
        },
        "hashes_salida": {
            key: sha256_file(path) for key, path in output_paths.items() if path.exists() and key != "manifest"
        },
    }


def print_console(
    dictamen_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    no_auto_df: pd.DataFrame,
    output_paths: Dict[str, Path],
    output_dir: Path,
    validator_state: str,
) -> None:
    print("AUDITORIA SIN REGISTROS 2025 CODCLI_LISTA 5809 — GENERADA")
    print("Fuentes originales modificadas: NO")
    print("Declaracion carga: NO_LISTO_PARA_CARGA")
    print(f"Estado validador: {validator_state}")
    print()
    print("DICTAMEN")
    print(dictamen_df.to_string(index=False))
    print()
    print("RESUMEN CAUSAS")
    print(
        resumen_df[
            ["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RESOLVER_AUTOMATICAMENTE"]
        ].to_string(index=False)
    )
    print()
    print("NO RESOLVER AUTOMATICO")
    print(no_auto_df[["CAUSA_PRINCIPAL_PROPUESTA", "FILA_5809"]].groupby("CAUSA_PRINCIPAL_PROPUESTA").count().reset_index().to_string(index=False))
    print()
    print("ARCHIVOS")
    print(f"Excel: {output_paths['excel']}")
    print(f"Informe: {output_paths['informe']}")
    print(f"Manifest: {output_paths['manifest']}")
    print(f"Script: {output_paths['script']}")
    print(f"Carpeta hito: {output_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Auditoria hito 56 sin registros 2025 CODCLI_LISTA 5809")
    parser.add_argument("--timestamp", default=None, help="Timestamp YYYYMMDD_HHMMSS. Si se omite, usa fecha/hora actual.")
    args = parser.parse_args()

    timestamp = args.timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = BASE_56 / f"AUDITORIA_SIN_REGISTROS_2025_CODCLI_LISTA_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)

    output_paths = {
        "excel": output_dir / OUTPUT_FILES["excel"],
        "informe": output_dir / OUTPUT_FILES["informe"],
        "manifest": output_dir / OUTPUT_FILES["manifest"],
        "script": output_dir / OUTPUT_FILES["script"],
    }

    require_paths()
    current_script = Path(__file__).resolve()
    target_script = output_paths["script"].resolve()
    if current_script != target_script:
        shutil.copy2(current_script, target_script)

    validator = run_validator_52(output_dir, timestamp)

    matriz_05 = read_sheet(MATRIZ_53, "05_SIN_REGISTROS_2025")
    cases = matriz_05[matriz_05["DICTAMEN_DICCIONARIO_ESTADO"] == "SIN_REGISTROS_2025_PARA_CODCLI_LISTA"].copy()
    if len(cases) != 122:
        raise RuntimeError(f"Se esperaban 122 casos SIN_REGISTROS_2025_PARA_CODCLI_LISTA; encontrados: {len(cases)}")

    hoja1, datos, matricula = prepare_promedios()
    mapeo = read_mapeo_for_cases(cases)
    case_tables = build_case_tables(cases, hoja1, datos, matricula, mapeo)
    casos_df = case_tables["casos"]
    resumen_df = build_resumen_causas(casos_df)
    causas_df = build_causas_propuestas(resumen_df)
    validacion_df = build_validation_df(validator)

    dictamen_df = pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Estado de carga": DECLARACION_CARGA,
                "Estado validador": validator["estado_validador"],
                "Fuentes originales modificadas": "NO",
                "Correcciones aplicadas": "NO",
                "SIES_READY generado": "NO",
                "Recalculo 20-21": "NO",
                "Mezcla hitos 54/55": "NO",
                "Casos fuente hito 53": len(cases),
                "Total casos auditados": len(casos_df),
                "Casos sin registros 2025 CODCLI_LISTA": int((casos_df["FILAS_2025_CODCLI_LISTA"] == 0).sum()),
                "Casos con RUT otro CODCLI 2025": int(
                    (casos_df["CODCLI_RUT_2025_DISTINTO_LISTA"].map(clean) != "").sum()
                ),
                "Casos con CODCLI_LISTA multiple/no exacto": int(
                    (casos_df["MAPEO_CODCLI_LISTA_MULTIPLE_O_NO_EXACTO"] == "SI").sum()
                ),
                "Dictamen global": (
                    "NO_LISTO_PARA_CARGA: auditoria diagnostica de 122 casos sin registros 2025; "
                    "no procede correccion automatica ni carga."
                ),
                "Proximo paso permitido": (
                    "Revision funcional y/o fuente academica complementaria; luego revalidar gobernanza antes de cualquier calculo."
                ),
            }
        ]
    )

    no_auto_df = case_tables["no_auto"].copy()
    sheets = OrderedDict(
        [
            ("00_DICTAMEN", dictamen_df),
            ("01_VALIDACION_GOBERNANZA", validacion_df),
            ("02_RESUMEN_CAUSAS", resumen_df),
            ("03_122_CASOS", casos_df),
            ("04_BUSQUEDA_CODCLI_OTROS_ANOS", case_tables["codcli_otros_anos"]),
            ("05_BUSQUEDA_RUT_PROMEDIOS", case_tables["rut_promedios"]),
            ("06_ESTADO_ACADEMICO_OBSERVADO", case_tables["estado_observado"]),
            ("07_MAPEO_CODCLI_LISTA", mapeo),
            ("08_CAUSAS_PROPUESTAS", causas_df),
            ("09_NO_RESOLVER_AUTOMATICO", no_auto_df),
        ]
    )

    # 10_FUENTES se agrega al final, cuando los productos ya existen o tienen ruta estable.
    placeholder_fuentes = build_fuentes_df(output_dir, output_paths, validator)
    sheets["10_FUENTES"] = placeholder_fuentes
    write_excel(output_paths["excel"], sheets)

    manifest = build_manifest(timestamp, output_dir, output_paths, validator, dictamen_df, resumen_df, sheets)
    md_text = build_markdown(output_dir, dictamen_df, validacion_df, resumen_df, casos_df, manifest)
    output_paths["informe"].write_text(md_text, encoding="utf-8")

    fuentes_df = build_fuentes_df(output_dir, output_paths, validator)
    sheets["10_FUENTES"] = fuentes_df
    write_excel(output_paths["excel"], sheets)

    manifest = build_manifest(timestamp, output_dir, output_paths, validator, dictamen_df, resumen_df, sheets)
    output_paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print_console(dictamen_df, resumen_df, no_auto_df, output_paths, output_dir, validator["estado_validador"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
