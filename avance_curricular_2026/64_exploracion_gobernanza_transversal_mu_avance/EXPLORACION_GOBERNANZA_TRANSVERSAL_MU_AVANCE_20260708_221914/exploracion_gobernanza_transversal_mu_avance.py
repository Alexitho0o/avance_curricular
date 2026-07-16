#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Hito exploratorio de gobernanza transversal entre Matricula Unificada 2026
y Avance Curricular SIES 2026.

No modifica fuentes originales, no genera carga, no genera SIES_READY, no
corrige datos y no aplica reglas transversales. Solo inventaria, compara y
clasifica respaldo documental.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO_PRINCIPAL = "Exploracion gobernanza transversal MU / Avance"
PROCESOS_COMPARADOS = "Matricula Unificada 2026 | Avance Curricular SIES 2026"
ANIO_MU = 2026
ANIO_AVANCE = 2026
ESTADO = "EXPLORATORIO"
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_64 = REPO / "avance_curricular_2026" / "64_exploracion_gobernanza_transversal_mu_avance"

MANUAL_MU_PDF = REPO / "Manual_Matrícula_Unificada_2026.pdf"
MANUAL_MU_TXT = REPO / "manual_matrícula_unificada.txt"
MANUAL_MU_DOCS = REPO / "docs" / "Manual_Matricula_Unificada_2026.pdf"
INSTRUCTIVO_AVANCE = (
    REPO
    / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/"
    / "Instructivo_Avance Curricular SIES - 2026.txt"
)
GOBERNANZA_51 = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)
VALIDADOR_52 = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx"
)

TARGET_TSVS = [
    REPO / "DURACION_ESTUDIOS.tsv",
    REPO / "gobernanza_escala_notas.tsv",
    REPO / "gobernanza_nac.tsv",
    REPO / "gobernanza_niveles.tsv",
    REPO / "gobernanza_pais_est_sec.tsv",
    REPO / "control/for_ing_act_trace_long.tsv",
    REPO / "control/gob_codcarpr_anioingreso_long.tsv",
    REPO / "REGLAS_MU2026.tsv",
    REPO / "gobernanza_for_ing_act.tsv",
]

OUTPUT_FILES = {
    "excel": "EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE.xlsx",
    "informe": "INFORME_EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE.md",
    "manifest": "manifest_exploracion_gobernanza_transversal_mu_avance.json",
    "script": "exploracion_gobernanza_transversal_mu_avance.py",
}

KEYWORDS = [
    "gobernanza",
    "control",
    "catálogo",
    "catalogo",
    "diccionario",
    "trace",
    "mapping",
    "mapeo",
    "regla",
    "nivel",
    "nacionalidad",
    "pais",
    "país",
    "duracion",
    "duración",
    "escala",
    "ingreso",
    "codigo",
    "código",
    "codcarpr",
    "codigo_unico",
    "plan_estudios",
    "manual",
    "instructivo",
]
ALLOWED_SUFFIXES = {".tsv", ".csv", ".xlsx", ".json", ".md", ".txt", ".pdf", ".py"}
EXCLUDE_DIRS = {".git", ".venv", "__pycache__", ".mypy_cache", ".pytest_cache"}

SHEETS = OrderedDict(
    [
        ("00_DICTAMEN_GLOBAL", "00_DICTAMEN_GLOBAL"),
        ("01_ARCHIVOS_EXPLORADOS", "01_ARCHIVOS_EXPLORADOS"),
        ("02_MANUALES_OFICIALES", "02_MANUALES_OFICIALES"),
        ("03_TSV_GOBERNANZA_LOCAL", "03_TSV_GOBERNANZA_LOCAL"),
        ("04_CRITERIOS_TRANSVERSALES_CANDIDATOS", "04_CRITERIOS_TRANSVERSALES"),
        ("05_CRITERIOS_NO_REUTILIZAR", "05_CRITERIOS_NO_REUTILIZAR"),
        ("06_COLUMNAS_COMPARABLES_MU_AVANCE", "06_COLUMNAS_COMPARABLES"),
        ("07_DICCIONARIOS_Y_CATALOGOS", "07_DICCIONARIOS_CATALOGOS"),
        ("08_PERIODO_ACADEMICO", "08_PERIODO_ACADEMICO"),
        ("09_NIVEL", "09_NIVEL"),
        ("10_DURACION_ESTUDIOS", "10_DURACION_ESTUDIOS"),
        ("11_ESCALA_NOTAS", "11_ESCALA_NOTAS"),
        ("12_NACIONALIDAD_PAIS_NIVELES", "12_NAC_PAIS_NIVELES"),
        ("13_CONTROL_TRACE", "13_CONTROL_TRACE"),
        ("14_BLOQUEOS_Y_PENDIENTES", "14_BLOQUEOS_PENDIENTES"),
        ("15_RECOMENDACION_OPERATIVA", "15_RECOMENDACION_OPERATIVA"),
        ("16_FUENTES", "16_FUENTES"),
        ("17_MANIFEST_LEGIBLE", "17_MANIFEST_LEGIBLE"),
    ]
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.lower()


def display(value: Any, max_len: int = 180) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_under(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def safe_read_tsv(path: Path) -> pd.DataFrame:
    try:
        return pd.read_csv(path, sep="\t", dtype=str).fillna("")
    except UnicodeDecodeError:
        return pd.read_csv(path, sep="\t", dtype=str, encoding="latin1").fillna("")
    except Exception:
        return pd.DataFrame()


def safe_read_text(path: Path) -> str:
    for encoding in ["utf-8", "utf-8-sig", "latin1"]:
        try:
            return path.read_text(encoding=encoding)
        except Exception:
            continue
    return ""


def classify_process(path: Path) -> str:
    text = normalize(path)
    if "matricula_unificada" in text or "matricula unificada" in text or "mu2026" in text or "manual_matricula" in text:
        return "Matricula Unificada 2026"
    if "avance_curricular" in text or "avance curricular" in text or "5809" in text or "5810" in text:
        return "Avance Curricular SIES 2026"
    if "extranjeros" in text or "nacionalidad" in text:
        return "Estudiantes Extranjeros / transversal posible"
    if "gobernanza" in text or "control" in text:
        return "Transversal o institucional"
    return "No determinado"


def classify_file(path: Path, exists: bool) -> Tuple[str, str, str, str, str, str]:
    if not exists:
        return ("pendiente", "archivo faltante", "E", "NO", "NO", "NO")
    name = normalize(path.name)
    full = normalize(path)
    if path.is_dir():
        return ("carpeta", "archivo de trabajo", "E", "NO", "NO", "NO")
    if "manual" in name or "instructivo" in name:
        return ("manual/instructivo", "manual oficial" if path.suffix.lower() in {".pdf", ".txt"} else "evidencia", "A", "SI", "NO", "NO")
    if path.suffix.lower() == ".py":
        return ("codigo", "codigo", "C", "NO", "NO", "SI")
    if "control" in full or "trace" in full:
        return ("control/auditoria", "auditoria", "C", "NO", "SI", "SI")
    if "gobernanza" in full or "regla" in full or "diccionario" in full or "catalogo" in full or "catálogo" in full:
        return ("gobernanza local", "decision interna", "D", "NO", "SI", "SI")
    if path.suffix.lower() in {".csv", ".xlsx", ".tsv"}:
        return ("dato tabular", "dato observado", "B", "NO", "SI", "NO")
    if path.suffix.lower() in {".json", ".md", ".txt"}:
        return ("documento/reporte", "evidencia", "E", "NO", "SI", "NO")
    return ("otro", "evidencia", "E", "NO", "NO", "NO")


def find_candidate_files() -> List[Path]:
    found: List[Path] = []
    for root, dirs, files in os.walk(REPO):
        root_path = Path(root)
        if is_under(root_path, BASE_64):
            dirs[:] = []
            continue
        dirs[:] = [d for d in dirs if d not in EXCLUDE_DIRS and not is_under(root_path / d, BASE_64)]
        for file in files:
            path = Path(root) / file
            if path.suffix.lower() not in ALLOWED_SUFFIXES:
                continue
            text = normalize(path.relative_to(REPO))
            if any(keyword in text for keyword in [normalize(k) for k in KEYWORDS]):
                found.append(path)
    for path in TARGET_TSVS + [MANUAL_MU_PDF, MANUAL_MU_TXT, MANUAL_MU_DOCS, INSTRUCTIVO_AVANCE, GOBERNANZA_51, VALIDADOR_52]:
        if path not in found:
            found.append(path)
    return sorted(set(found), key=lambda p: str(p).lower())


def build_archivos_explorados() -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    folders = [
        REPO,
        REPO / "avance_curricular_2026",
        REPO / "control",
        REPO / "resultados",
        REPO / "gobernanza_columnas_mu",
    ]
    for i in range(51, 64):
        matches = sorted((REPO / "avance_curricular_2026").glob(f"{i}_*"))
        folders.extend(matches)
    for path in folders + find_candidate_files():
        exists = path.exists()
        tipo, clasif, nivel, gov_off, gov_data, gov_tech = classify_file(path, exists)
        rows.append(
            {
                "Ruta": str(path),
                "Existe SI/NO": "SI" if exists else "NO",
                "Tipo": tipo,
                "Clasificacion": clasif,
                "Proceso probable": classify_process(path),
                "Subproyecto probable": infer_subproject(path),
                "Nivel de respaldo A/B/C/D/E": nivel,
                "Puede gobernar reglas oficiales SI/NO": gov_off,
                "Puede gobernar datos observados SI/NO": gov_data,
                "Puede gobernar transformacion tecnica SI/NO": gov_tech,
                "Hash SHA256 si existe": sha256_file(path) if exists and path.is_file() else "",
                "Observacion": observation_for_file(path, exists, nivel),
            }
        )
    return pd.DataFrame(rows).drop_duplicates(subset=["Ruta"]).sort_values(["Existe SI/NO", "Ruta"], ascending=[False, True])


def infer_subproject(path: Path) -> str:
    text = normalize(path)
    if "51_gobernanza" in text:
        return "Hito 51 gobernanza integral Avance 5809"
    if "52_validador" in text:
        return "Hito 52 validador precalculo"
    if "58_consolidado" in text:
        return "Hito 58 consolidado resolucion 16-19"
    if "duracion" in text:
        return "Duracion de estudios / carreras"
    if "nivel" in text:
        return "Nivel carrera/global"
    if "nac" in text or "pais" in text or "país" in text:
        return "Nacionalidad/pais"
    if "escala" in text or "nota" in text:
        return "Escala de notas"
    if "for_ing_act" in text or "ingreso" in text:
        return "Forma/anio ingreso"
    if "codcarpr" in text:
        return "CODCARPR / oferta academica"
    return "Exploracion transversal"


def observation_for_file(path: Path, exists: bool, nivel: str) -> str:
    if not exists:
        return "Archivo/carpeta esperado no encontrado localmente."
    if nivel == "A":
        return "Fuente oficial solo para su propio proceso y periodo; no habilita reutilizacion automatica."
    if nivel == "D":
        return "Gobernanza local o decision institucional; no reemplaza manual/instructivo."
    if nivel == "C":
        return "Implementacion/auditoria tecnica; no es fuente normativa."
    if nivel == "B":
        return "Dato observado; puede respaldar evidencia, no norma."
    return "Evidencia exploratoria o pendiente."


def extract_manual_rows() -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    sources = [
        ("Matricula Unificada 2026", MANUAL_MU_PDF, MANUAL_MU_TXT),
        ("Avance Curricular SIES 2026", INSTRUCTIVO_AVANCE, INSTRUCTIVO_AVANCE),
    ]
    topics = {
        "nivel": ["nivel", "pregrado", "postgrado", "postitulo"],
        "duracion": ["duración", "duracion", "semestres"],
        "periodo_semestre": ["semestre", "periodo", "1er", "2º", "2do"],
        "codigo_plan": ["codigo_unico", "código único", "plan_estudios", "plan de estudios", "codigo sies"],
        "nacionalidad_pais": ["nacionalidad", "país", "pais"],
        "escala_notas": ["nota", "escala", "promedio"],
        "avance_unidades": ["unidades cursadas", "unidades aprobadas", "avance curricular", "convalid", "homolog"],
        "vigencia": ["vigencia"],
    }
    for proceso, official_path, text_path in sources:
        text = safe_read_text(text_path)
        lines = text.splitlines()
        for topic, terms in topics.items():
            matches = []
            for idx, line in enumerate(lines, start=1):
                low = normalize(line)
                if any(normalize(term) in low for term in terms):
                    matches.append((idx, clean(line)))
                if len(matches) >= 8:
                    break
            if matches:
                snippet = " | ".join(f"L{idx}: {line}" for idx, line in matches[:4])
                rows.append(
                    {
                        "Proceso": proceso,
                        "Fuente": str(official_path),
                        "Seccion o referencia si esta disponible": topic,
                        "Regla textual o resumen trazable": snippet,
                        "Campo/tema": topic,
                        "Aplica a": proceso,
                        "Nivel respaldo A": "A",
                        "Observacion": (
                            "Extracto/resumen local. La regla oficial aplica solo al proceso fuente; "
                            "requiere confirmacion para uso transversal."
                        ),
                    }
                )
            else:
                rows.append(
                    {
                        "Proceso": proceso,
                        "Fuente": str(official_path),
                        "Seccion o referencia si esta disponible": topic,
                        "Regla textual o resumen trazable": "No se encontro evidencia local por busqueda de terminos.",
                        "Campo/tema": topic,
                        "Aplica a": proceso,
                        "Nivel respaldo A": "A",
                        "Observacion": "Ausencia de evidencia textual local no equivale a negacion; queda pendiente.",
                    }
                )
    return pd.DataFrame(rows)


def summarize_tsv(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {
            "Archivo": str(path),
            "Columnas": "",
            "Filas": 0,
            "Campos gobernados": "",
            "Valores principales": "",
            "Uso probable": "PENDIENTE_ARCHIVO_NO_EXISTE",
            "Proceso probable": classify_process(path),
            "Nivel respaldo": "E",
            "Riesgo": "No existe localmente.",
            "Observacion": "Archivo señalado no encontrado.",
        }
    df = safe_read_tsv(path)
    columns = list(df.columns)
    sample_values = []
    for col in columns[:5]:
        vals = sorted(set(clean(v) for v in df[col].head(20).tolist() if clean(v)))[:5]
        if vals:
            sample_values.append(f"{col}: {', '.join(vals)}")
    _, _, nivel, _, _, _ = classify_file(path, True)
    return {
        "Archivo": str(path),
        "Columnas": " | ".join(columns),
        "Filas": len(df),
        "Campos gobernados": infer_governed_fields(path, columns),
        "Valores principales": " || ".join(sample_values),
        "Uso probable": infer_tsv_use(path),
        "Proceso probable": classify_process(path),
        "Nivel respaldo": nivel,
        "Riesgo": risk_for_tsv(path),
        "Observacion": "TSV local inventariado; no sustituye manual/instructivo oficial.",
    }


def infer_governed_fields(path: Path, columns: Iterable[str]) -> str:
    text = normalize(path.name + " " + " ".join(columns))
    fields = []
    for key in ["CODIGO_UNICO", "PLAN_ESTUDIOS", "CODCARPR", "DURACION_ESTUDIOS", "NIVEL", "NACIONALIDAD", "PAIS", "ESCALA", "NOTA", "FOR_ING_ACT", "PERIODO", "ANOINGRESO"]:
        if normalize(key) in text:
            fields.append(key)
    return " | ".join(fields) if fields else "Campos a determinar por revision funcional"


def infer_tsv_use(path: Path) -> str:
    text = normalize(path)
    if "escala" in text:
        return "Conversion escala de notas para Matricula Unificada."
    if "nac" in text:
        return "Catalogo nacionalidad/pais para Matricula Unificada o Extranjeros."
    if "pais_est_sec" in text:
        return "Catalogo pais establecimiento secundario."
    if "niveles" in text or "nivel" in text:
        return "Catalogo nivel global/carrera o auditoria de niveles."
    if "duracion" in text:
        return "Duracion de estudios/titulacion/total por codigo unico/carrera."
    if "for_ing_act" in text:
        return "Forma de ingreso actual y trazas de decision MU."
    if "codcarpr" in text:
        return "Control CODCARPR, jornada y anio ingreso."
    return "Uso probable pendiente."


def risk_for_tsv(path: Path) -> str:
    text = normalize(path)
    if "escala" in text:
        return "Riesgo alto si se reutiliza en Avance: columnas 16-19 no se gobiernan por notas."
    if "for_ing_act" in text or "control" in text or "trace" in text:
        return "Riesgo alto: archivo tecnico/de auditoria, no norma oficial."
    return "Riesgo medio: requiere confirmar alcance por proceso antes de reutilizar."


def build_tsv_inventory() -> pd.DataFrame:
    tsvs = set(TARGET_TSVS)
    for path in find_candidate_files():
        if path.suffix.lower() == ".tsv":
            tsvs.add(path)
    rows = [summarize_tsv(path) for path in sorted(tsvs, key=lambda p: str(p).lower())]
    return pd.DataFrame(rows)


def build_candidates() -> pd.DataFrame:
    rows = [
        candidate("NIVEL_CARRERA / NIVEL_GLOBAL", "Catalogo de niveles de carrera/global aparece en MU y Avance, pero el uso operativo cambia por proceso.", "Manual MU + gobernanza_niveles.tsv", "Instructivo Avance Anexo carreras + hito 51", "gobernanza_niveles.tsv", "SI", "SI", "PENDIENTE", "A/D", "SI", "Usar nivel MU para Avance sin validar puede clasificar mal carreras/programas.", "CANDIDATO_CON_VALIDACION"),
        candidate("DURACION_ESTUDIOS", "Duracion por codigo unico/plan/carrera existe en Avance y TSV local; puede compartir universo de carreras, no regla de calculo.", "Manual MU menciona duracion minima y validaciones de trayectoria.", "Instructivo Avance define duracion estudios/titulacion/total.", "DURACION_ESTUDIOS.tsv", "PENDIENTE", "SI", "PENDIENTE", "A/D", "SI", "Confundir duracion de estudios con regla de matricula o avance puede afectar validaciones.", "CANDIDATO_CON_VALIDACION"),
        candidate("CODIGO_UNICO + PLAN_ESTUDIOS", "Identificadores de programa/carrera aparecen en ambos procesos.", "Manual MU / resultados MU", "Instructivo Avance y precargas 5809/5810", "DURACION_ESTUDIOS.tsv / hito 51", "SI", "SI", "PENDIENTE", "A/B/D", "SI", "Misma llave aparente no implica misma cardinalidad ni regla.", "CANDIDATO_CON_VALIDACION"),
        candidate("VIGENCIA", "Campo existe en ambos procesos, pero el significado de mantener/eliminar puede ser especifico.", "Manual MU", "Instructivo Avance", "hito 51", "SI", "SI", "PENDIENTE", "A", "SI", "Reutilizar codigos sin validar contexto puede eliminar registros indebidamente.", "CANDIDATO_BLOQUEADO_HASTA_CONFIRMACION"),
        candidate("NACIONALIDAD / PAIS", "Catalogos locales existen para MU/extranjeros; no son parte central de Avance 5809 16-19.", "Manual MU Cuadro pais/nacionalidad + TSV", "Sin evidencia de uso en Avance 5809 16-19.", "gobernanza_nac.tsv / gobernanza_pais_est_sec.tsv", "SI", "NO", "NO", "A/D", "SI", "Aplicar en Avance agregaria criterio ajeno al proceso.", "NO_REUTILIZAR_EN_AVANCE_5809"),
        candidate("ESCALA_NOTAS", "Conversion de notas esta gobernada localmente para MU; Avance 16-19 usa unidades/estados, no notas.", "Manual MU Anexo 7 + gobernanza_escala_notas.tsv", "Sin regla de notas para 16-19; hito 51 gobierna ESTADO/DESCRIPCION_ESTADO.", "gobernanza_escala_notas.tsv", "SI", "NO", "NO", "A/D", "SI", "Usar notas para resolver Avance contradice diccionario de estados y objetivo 16-19.", "NO_REUTILIZAR_EN_AVANCE"),
        candidate("PERIODO academico raw", "PERIODO/PERIODOMATRICULA aparecen en datos; mapping a semestres requiere fuente especifica.", "Manual MU usa semestre ingreso 1/2; trazas tienen PERIODOMATRICULA.", "Instructivo Avance pide 1er/2do semestre, pero PROMEDIOS PERIODO tiene 1..5.", "control trace / PROMEDIOS / hito 55", "PENDIENTE", "PENDIENTE", "PENDIENTE", "B/C/D", "SI", "Asumir 2/3=segundo semestre sin fuente puede cambiar CURSO_2DO_SEM.", "PENDIENTE_BLOQUEADO"),
        candidate("FOR_ING_ACT", "Forma de ingreso actual esta gobernada para MU, no para Avance 5809.", "Manual MU + gobernanza_for_ing_act.tsv + REGLAS_MU2026.tsv", "Sin evidencia oficial para Avance.", "control/for_ing_act_trace_long.tsv", "SI", "NO", "NO", "A/D/C", "SI", "Reciclar reglas MU de ingreso en Avance agregaria norma externa.", "NO_REUTILIZAR_EN_AVANCE"),
        candidate("ESTADO/DESCRIPCION_ESTADO PROMEDIOS", "Diccionario de estados academicos usado en Avance 16-19; no es regla oficial MU.", "No aplica como regla MU.", "Hito 51 y PROMEDIOS: A/E/I/R/NULL.", "GOBERNANZA_51", "NO", "SI", "NO", "B/D", "SI", "Usar diccionario de Avance en MU sin manual generaria criterio no oficial.", "NO_REUTILIZAR_EN_MU"),
    ]
    return pd.DataFrame(rows)


def candidate(criterio: str, desc: str, mu: str, avance: str, local: str, aplica_mu: str, aplica_av: str, trans: str, nivel: str, valida: str, riesgo: str, dictamen: str) -> Dict[str, Any]:
    return {
        "Criterio": criterio,
        "Descripcion": desc,
        "Fuente Matricula Unificada": mu,
        "Fuente Avance Curricular": avance,
        "Fuente local institucional": local,
        "Aplica a MU": aplica_mu,
        "Aplica a Avance": aplica_av,
        "Aplica transversalmente": trans,
        "Nivel de respaldo": nivel,
        "Requiere validacion funcional": valida,
        "Riesgo si se reutiliza sin validar": riesgo,
        "Dictamen": dictamen,
    }


def build_no_reuse() -> pd.DataFrame:
    rows = [
        ("Escala de notas MU", "Matricula Unificada", "Avance Curricular 5809", "Avance 16-19 se gobierna por unidades/estados, no conversion de notas.", "ALTO", "gobernanza_escala_notas.tsv / Manual MU", "BLOQUEADO"),
        ("FOR_ING_ACT", "Matricula Unificada", "Avance Curricular", "Campo/reglas de forma de ingreso MU no aparecen como regla Avance.", "ALTO", "gobernanza_for_ing_act.tsv / REGLAS_MU2026.tsv", "BLOQUEADO"),
        ("Nacionalidad/Pais establecimiento", "Matricula Unificada / Extranjeros", "Avance Curricular 5809", "No gobierna columnas 16-19 ni avance anual/acumulado.", "MEDIO", "gobernanza_nac.tsv / gobernanza_pais_est_sec.tsv", "NO_REUTILIZAR"),
        ("PERIODO 2/3 como segundo semestre", "Decision tecnica previa", "Ambos procesos", "No hay evidencia oficial suficiente para homologar PERIODO raw 2/3 a semestre.", "ALTO", "hito 55 / control trace", "PENDIENTE_BLOQUEADO"),
        ("N_CODCLI como CODCLI academico", "Implementacion erronea previa", "Avance Curricular", "Hito 51 lo bloquea expresamente; es conteo, no llave.", "CRITICO", "GOBERNANZA_51", "BLOQUEADO"),
        ("PES_READY MU", "Matricula Unificada", "Avance Curricular", "Archivo de carga de otro proceso no gobierna Avance ni este hito.", "CRITICO", "resultados/*PES_READY*", "NO_REUTILIZAR"),
        ("Diccionario ESTADO PROMEDIOS Avance", "Avance Curricular", "Matricula Unificada", "Dato observado/gobernanza Avance; no reemplaza manual MU.", "MEDIO", "GOBERNANZA_51 / PROMEDIOS", "NO_REUTILIZAR_SIN_MANUAL"),
    ]
    return pd.DataFrame(rows, columns=["Criterio", "Proceso origen", "Proceso donde NO aplicar", "Motivo", "Riesgo", "Fuente", "Estado"])


def build_columnas_comparables() -> pd.DataFrame:
    rows = [
        ("Codigo carrera/programa", "CODIGO_UNICO / CODIGO_SIES", "CODIGO_UNICO", "SI", "PENDIENTE", "PENDIENTE", "Misma finalidad de identificacion; reglas/cardinalidad dependen del proceso."),
        ("Plan de estudios", "Plan/carrera actual segun MU", "PLAN_ESTUDIOS", "PARCIAL", "PENDIENTE", "PENDIENTE", "En Avance es llave oficial con CODIGO_UNICO; en MU puede estar implícito en oferta."),
        ("Nivel carrera", "NIVEL / NIVEL_CARRERA", "NIVEL_CARRERA", "PARCIAL", "PENDIENTE", "NO", "Catalogo comparable; no usar sin confirmar alcance."),
        ("Duracion estudios", "Duracion programa/carrera", "DURACION_ESTUDIOS/TITULACION/TOTAL", "PARCIAL", "PENDIENTE", "NO", "Avance define distribucion de unidades por duracion; MU usa criterios de matricula/trayectoria."),
        ("Semestre ingreso", "SEM_ING_ACT/ORI", "SEM_INGRESO_CARRERA_ACTUAL/ORIGEN", "SI", "PENDIENTE", "PENDIENTE", "Ambos usan semestre 1/2 para ingreso, pero no equivale a PERIODO raw de asignaturas."),
        ("Vigencia", "VIGENCIA", "VIGENCIA", "SI", "PENDIENTE", "PENDIENTE", "Mismo nombre no asegura misma regla operacional."),
        ("Nacionalidad", "NACIONALIDAD/COD_NAC", "No central en 5809", "NO", "NO", "NO", "No reutilizar para Avance 16-19."),
        ("Notas/promedios", "NOTA/PROMEDIO", "No usado para 16-19", "NO", "NO", "NO", "Avance 16-19 se diagnostica por unidades y estados."),
    ]
    return pd.DataFrame(rows, columns=["Campo o concepto", "Columna en MU si existe", "Columna en Avance si existe", "Misma finalidad SI/NO/PARCIAL", "Misma regla SI/NO/PENDIENTE", "Misma fuente SI/NO", "Observacion"])


def build_diccionarios_catalogos() -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for path in TARGET_TSVS:
        if not path.exists():
            continue
        df = safe_read_tsv(path)
        if df.empty:
            continue
        code_col, desc_col = infer_code_desc_cols(df)
        for _, row in df.head(30).iterrows():
            rows.append(
                {
                    "Archivo": str(path),
                    "Hoja o TSV": path.name,
                    "Codigo": clean(row.get(code_col, "")) if code_col else "",
                    "Descripcion": clean(row.get(desc_col, "")) if desc_col else "",
                    "Campo gobernado": infer_governed_fields(path, df.columns),
                    "Proceso": classify_process(path),
                    "Uso permitido": infer_tsv_use(path),
                    "Uso prohibido": "No usar como regla oficial transversal sin manual del proceso receptor.",
                    "Nivel respaldo": classify_file(path, True)[2],
                    "Observacion": "Muestra de catalogo/diccionario local.",
                }
            )
    if GOBERNANZA_51.exists():
        try:
            dic = pd.read_excel(GOBERNANZA_51, sheet_name="05_DICCIONARIO_ESTADO_PROMEDIOS", dtype=str).fillna("")
            for _, row in dic.iterrows():
                rows.append(
                    {
                        "Archivo": str(GOBERNANZA_51),
                        "Hoja o TSV": "05_DICCIONARIO_ESTADO_PROMEDIOS",
                        "Codigo": clean(row.get("Código", row.get("Codigo", row.get("ESTADO", "")))),
                        "Descripcion": clean(row.get("Descripción", row.get("Descripcion", row.get("DESCRIPCION_ESTADO", "")))),
                        "Campo gobernado": "ESTADO / DESCRIPCION_ESTADO PROMEDIOS",
                        "Proceso": "Avance Curricular SIES 2026",
                        "Uso permitido": "Interpretacion gobernada para auditorias Avance 16-19.",
                        "Uso prohibido": "No usar como regla oficial MU ni como sustituto de instructivo.",
                        "Nivel respaldo": "D/B",
                        "Observacion": "Diccionario observado/gobernado en hito 51.",
                    }
                )
        except Exception:
            pass
    return pd.DataFrame(rows)


def infer_code_desc_cols(df: pd.DataFrame) -> Tuple[str, str]:
    cols = list(df.columns)
    code = next((c for c in cols if any(k in normalize(c) for k in ["cod", "codigo", "for_ing_act", "nivel", "nota"])), cols[0] if cols else "")
    desc = next((c for c in cols if any(k in normalize(c) for k in ["desc", "descripcion", "nombre", "regla"])), cols[1] if len(cols) > 1 else code)
    return code, desc


def build_periodo() -> pd.DataFrame:
    rows = []
    treatments = {
        "1": ("Aparece en SEM_INGRESO y en PERIODO/PERIODOMATRICULA observado.", "MU: semestre ingreso 1 oficial; trace MU observado.", "Avance: instructivo menciona 1er semestre; PERIODO raw observado en PROMEDIOS.", "PARCIAL", "A/B/C", "Usar como semestre solo cuando el campo oficial sea semestre; PERIODO raw requiere confirmacion."),
        "2": ("Aparece en SEM_INGRESO y PERIODO raw.", "MU: semestre ingreso 2 oficial; PERIODOMATRICULA observado.", "Avance: instructivo menciona 2do semestre; PERIODO raw observado.", "PARCIAL", "A/B/C", "No asumir que todo PERIODO=2 agota segundo semestre sin fuente."),
        "3": ("Aparece como PERIODO raw en PROMEDIOS/hitos, no como semestre oficial.", "No confirmado para MU como regla transversal.", "En hito 55 fue causa PERIODO_NO_1_2; requiere decision funcional.", "NO", "B/D/E", "No asumir 2/3=segundo semestre; queda bloqueado."),
        "4": ("Aparece como PERIODO raw observado.", "No confirmado.", "No confirmado para 16-17; hito 55 lo bloquea.", "NO", "B/D/E", "Bloqueado hasta fuente funcional."),
        "5": ("Aparece como PERIODO raw observado.", "No confirmado.", "No confirmado para 16-17; hito 55 lo bloquea.", "NO", "B/D/E", "Bloqueado hasta fuente funcional."),
    }
    for periodo, vals in treatments.items():
        rows.append(
            {
                "Periodo": periodo,
                "Fuente donde aparece": vals[0],
                "Tratamiento observado": vals[5],
                "Tratamiento en MU": vals[1],
                "Tratamiento en Avance": vals[2],
                "Si esta gobernado": vals[3],
                "Si es decision interna": "SI" if periodo in {"3", "4", "5"} else "PENDIENTE",
                "Si queda pendiente": "SI",
                "Nivel respaldo": vals[4],
                "Dictamen": "NO_APLICAR_TRANSVERSALMENTE_SIN_CONFIRMACION",
            }
        )
    return pd.DataFrame(rows)


def build_nivel() -> pd.DataFrame:
    rows = []
    df = safe_read_tsv(REPO / "gobernanza_niveles.tsv") if (REPO / "gobernanza_niveles.tsv").exists() else pd.DataFrame()
    for _, row in df.iterrows():
        rows.append(
            {
                "Fuente local": str(REPO / "gobernanza_niveles.tsv"),
                "Manual MU": "Manual MU menciona niveles Pregrado/Postgrado/Postitulo y nivel academico.",
                "Instructivo Avance": "Anexo carreras contiene NIVEL_CARRERA y categorias.",
                "Archivo gobernanza_niveles.tsv": f"{clean(row.get('NIVEL_GLOBAL'))}-{clean(row.get('NIVEL_GLOBAL_DESC'))} / {clean(row.get('NIVEL_CARRERA'))}-{clean(row.get('NIVEL_CARRERA_DESC'))}",
                "Aplicacion a columnas": "NIVEL_GLOBAL, NIVEL_CARRERA, NIVEL",
                "Si aplica a ambos procesos": "PENDIENTE",
                "Si requiere validacion": "SI",
                "Riesgo": "Mismo catalogo aparente; regla de uso por proceso puede diferir.",
            }
        )
    if not rows:
        rows.append({"Fuente local": str(REPO / "gobernanza_niveles.tsv"), "Manual MU": "Pendiente", "Instructivo Avance": "Pendiente", "Archivo gobernanza_niveles.tsv": "No encontrado", "Aplicacion a columnas": "", "Si aplica a ambos procesos": "PENDIENTE", "Si requiere validacion": "SI", "Riesgo": "Sin archivo local."})
    return pd.DataFrame(rows)


def build_duracion() -> pd.DataFrame:
    path = REPO / "DURACION_ESTUDIOS.tsv"
    df = safe_read_tsv(path) if path.exists() else pd.DataFrame()
    rows = []
    if not df.empty:
        grouped = df.groupby(["TIPO_PLAN_CARRERA", "DURACION_ESTUDIOS"], dropna=False).size().reset_index(name="CASOS").head(40)
        for _, row in grouped.iterrows():
            rows.append(
                {
                    "Columnas": " | ".join(df.columns),
                    "Valores": f"TIPO_PLAN_CARRERA={clean(row.get('TIPO_PLAN_CARRERA'))}; DURACION_ESTUDIOS={clean(row.get('DURACION_ESTUDIOS'))}; CASOS={row.get('CASOS')}",
                    "Uso observado": "Catalogo local por CODIGO_UNICO/CODIGO_CARRERA/CODCARPR.",
                    "Relacion con tipo de plan": clean(row.get("TIPO_PLAN_CARRERA")),
                    "Relacion con codigo unico": "Incluye CODIGO_UNICO; no implica regla transversal automatica.",
                    "Aplicacion a MU": "PENDIENTE",
                    "Aplicacion a Avance": "SI para carreras Avance si coincide con instructivo y precarga.",
                    "Estado gobernanza": "DATO_LOCAL_GOBERNADO_PENDIENTE_TRANSVERSAL",
                }
            )
    else:
        rows.append({"Columnas": "", "Valores": "", "Uso observado": "Archivo no disponible", "Relacion con tipo de plan": "", "Relacion con codigo unico": "", "Aplicacion a MU": "PENDIENTE", "Aplicacion a Avance": "PENDIENTE", "Estado gobernanza": "PENDIENTE"})
    return pd.DataFrame(rows)


def build_escala() -> pd.DataFrame:
    path = REPO / "gobernanza_escala_notas.tsv"
    df = safe_read_tsv(path) if path.exists() else pd.DataFrame()
    rows = []
    if not df.empty:
        for _, row in df.head(30).iterrows():
            rows.append(
                {
                    "Archivo": str(path),
                    "NOTA_FUENTE": clean(row.get("NOTA_FUENTE", "")),
                    "NOTA_MU": clean(row.get("NOTA_MU", "")),
                    "Regla conversion": clean(row.get("REGLA_CONVERSION", "")),
                    "Si aplica a Avance Curricular": "NO",
                    "Si Avance usa notas o no": "NO para columnas 16-19; usa unidades y estados.",
                    "Si pertenece a otro proceso": "SI: Matricula Unificada.",
                    "Riesgo de reutilizacion indebida": "ALTO",
                    "Dictamen": "NO_REUTILIZAR_EN_AVANCE_16_19",
                }
            )
    else:
        rows.append({"Archivo": str(path), "NOTA_FUENTE": "", "NOTA_MU": "", "Regla conversion": "", "Si aplica a Avance Curricular": "PENDIENTE", "Si Avance usa notas o no": "No evaluado", "Si pertenece a otro proceso": "PENDIENTE", "Riesgo de reutilizacion indebida": "ALTO", "Dictamen": "PENDIENTE"})
    return pd.DataFrame(rows)


def build_nac_pais() -> pd.DataFrame:
    rows = []
    for path in [REPO / "gobernanza_nac.tsv", REPO / "gobernanza_pais_est_sec.tsv", REPO / "gobernanza_niveles.tsv"]:
        df = safe_read_tsv(path) if path.exists() else pd.DataFrame()
        rows.append(
            {
                "Archivo": str(path),
                "Columnas": " | ".join(df.columns) if not df.empty else "",
                "Filas": len(df),
                "Aplica a Matricula Unificada": "SI" if "nivel" not in normalize(path.name) else "PENDIENTE",
                "Aplica a Avance Curricular": "NO para 5809 16-19" if "nivel" not in normalize(path.name) else "PENDIENTE",
                "Aplica a Estudiantes Extranjeros": "SI/PENDIENTE" if "nac" in normalize(path.name) or "pais" in normalize(path.name) else "NO",
                "Aplica a Otro proceso": "PENDIENTE",
                "No mezclar procesos": "SI",
                "Dictamen": "NO_REUTILIZAR_TRANSVERSAL_SIN_MANUAL",
            }
        )
    return pd.DataFrame(rows)


def build_control_trace() -> pd.DataFrame:
    rows = []
    for path in [REPO / "control/for_ing_act_trace_long.tsv", REPO / "control/gob_codcarpr_anioingreso_long.tsv"]:
        df = safe_read_tsv(path) if path.exists() else pd.DataFrame()
        rows.append(
            {
                "Archivo": str(path),
                "Columnas afectadas": " | ".join(df.columns) if not df.empty else "",
                "Transformaciones que gobierna": infer_tsv_use(path),
                "Aplica a MU": "SI/PENDIENTE",
                "Aplica a Avance": "NO/PENDIENTE; no usar sin confirmacion documental.",
                "Tipo": "auditoria/implementacion/decision",
                "Nivel respaldo": "C/D",
                "Riesgo": "No es fuente normativa; alto riesgo si se convierte en regla transversal.",
                "Filas": len(df),
            }
        )
    return pd.DataFrame(rows)


def build_bloqueos() -> pd.DataFrame:
    rows = [
        ("PERIODO raw 3/4/5", "No confirmado como semestre para Avance ni MU.", "Validacion funcional/documental.", "ALTO", "BLOQUEADO"),
        ("Escala notas en Avance 16-19", "Pertenece a MU; Avance usa unidades/estados.", "No reutilizar.", "ALTO", "BLOQUEADO"),
        ("FOR_ING_ACT en Avance", "Regla MU no aplica automaticamente.", "Manual Avance o decision funcional formal.", "ALTO", "BLOQUEADO"),
        ("Nacionalidad/Pais en Avance 5809", "No gobierna columnas 16-19.", "No reutilizar salvo nuevo alcance.", "MEDIO", "BLOQUEADO"),
        ("DURACION_ESTUDIOS como regla MU", "TSV local no reemplaza manual MU.", "Confirmar uso exacto por proceso.", "MEDIO", "PENDIENTE"),
        ("NIVEL compartido", "Catalogo comparable, regla no cerrada transversalmente.", "Confirmacion funcional y manuales.", "MEDIO", "PENDIENTE"),
        ("CODIGO_UNICO/PLAN_ESTUDIOS transversal", "Misma finalidad, no necesariamente misma cardinalidad.", "Gobernanza de llave por proceso.", "ALTO", "PENDIENTE"),
        ("Criterios de hitos 60-63", "Son terminales/revisiones posteriores, no oficiales por si solos.", "Incorporar solo tras validacion funcional.", "MEDIO", "PENDIENTE"),
    ]
    return pd.DataFrame(rows, columns=["Criterio no confirmado", "Contradiccion o regla sin respaldo oficial", "Requisito para cerrar", "Riesgo", "Estado"])


def build_recomendacion() -> pd.DataFrame:
    rows = [
        ("Usar ahora", "Inventario y clasificacion de fuentes por respaldo.", "Permitido", "No aplica reglas de negocio."),
        ("Usar ahora", "Hito 51 y 52 para mantener bloqueos Avance 5809.", "Permitido", "Solo dentro de Avance y con alcance declarado."),
        ("No usar", "Escala notas MU para Avance 16-19.", "Bloqueado", "Proceso y campos distintos."),
        ("No usar", "FOR_ING_ACT MU para Avance.", "Bloqueado", "No hay respaldo oficial Avance."),
        ("No usar", "PERIODO 2/3=segundo semestre.", "Bloqueado", "No hay fuente oficial suficiente."),
        ("Requiere validacion", "Nivel carrera/global transversal.", "Pendiente", "Catalogo comparable, regla no cerrada."),
        ("Requiere validacion", "Duracion estudios como criterio compartido.", "Pendiente", "Debe separarse codigo unico/plan/carrera/matricula/conciliacion."),
        ("Incorporar a validador", "Bloqueo de reglas transversales no confirmadas.", "Recomendado", "Nuevo validador transversal o extension hito 52."),
        ("Siguiente hito", "Gobernanza transversal formal si responsables validan candidatos.", "Recomendado", "Convertir candidatos en matriz de decision con firmas/responsables."),
    ]
    return pd.DataFrame(rows, columns=["Categoria", "Criterio", "Recomendacion", "Observacion"])


def build_fuentes(paths_df: pd.DataFrame, output_paths: Dict[str, Path], desktop_dir: Path) -> pd.DataFrame:
    rows = []
    for _, row in paths_df.iterrows():
        if clean(row.get("Existe SI/NO")) == "SI":
            rows.append(
                {
                    "Tipo": "INPUT",
                    "Ruta": row["Ruta"],
                    "Hash SHA256": row["Hash SHA256 si existe"],
                    "Nivel respaldo": row["Nivel de respaldo A/B/C/D/E"],
                    "Uso": row["Clasificacion"],
                }
            )
    for key, path in output_paths.items():
        rows.append(
            {
                "Tipo": "OUTPUT",
                "Ruta": str(path),
                "Hash SHA256": sha256_file(path) if path.exists() and key != "manifest" else "",
                "Nivel respaldo": "DERIVADO",
                "Uso": "Producto exploratorio hito 64.",
            }
        )
    rows.append({"Tipo": "OUTPUT_DESKTOP", "Ruta": str(desktop_dir), "Hash SHA256": "", "Nivel respaldo": "DERIVADO", "Uso": "Copia de productos al Escritorio."})
    return pd.DataFrame(rows)


def build_manifest(timestamp: str, output_dir: Path, desktop_dir: Path, output_paths: Dict[str, Path], sheets: Dict[str, pd.DataFrame], counts: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "proceso_principal": PROCESO_PRINCIPAL,
        "procesos_comparados": PROCESOS_COMPARADOS,
        "timestamp": timestamp,
        "estado": ESTADO,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "archivos_carga_generados": "NO",
        "sies_ready_generado": "NO",
        "correcciones_aplicadas": "NO",
        "reglas_transversales_aplicadas": "NO",
        "genera_archivo_carga": "NO",
        "genera_sies_ready": "NO",
        "corrige_datos": "NO",
        "aplica_reglas_transversales": "NO",
        "carpeta_salida": str(output_dir),
        "carpeta_escritorio": str(desktop_dir),
        "archivos": {key: str(path) for key, path in output_paths.items()},
        "productos": {key: str(path) for key, path in output_paths.items()},
        "conteos": counts,
        "hojas_excel": {name: len(df) for name, df in sheets.items()},
        "hashes_salida": {key: sha256_file(path) for key, path in output_paths.items() if path.exists() and key != "manifest"},
    }


def manifest_legible(manifest: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    for key, value in manifest.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False)
        rows.append({"Clave": key, "Valor": value})
    return pd.DataFrame(rows)


def excel_safe_frame(df: pd.DataFrame) -> pd.DataFrame:
    def safe_value(value: Any) -> Any:
        if isinstance(value, str) and len(value) > 32000:
            return value[:31980] + " [...TRUNCADO_PARA_EXCEL]"
        return value

    return df.map(safe_value)


def write_excel(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for logical_name, df in sheets.items():
            sheet_name = SHEETS.get(logical_name, logical_name)[:31]
            excel_safe_frame(df).to_excel(writer, sheet_name=sheet_name, index=False)
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            values = [display(cell.value, 80) for cell in column_cells[:150]]
            width = min(max([len(v) for v in values] + [10]) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    wb.save(path)


def build_markdown(dictamen: pd.DataFrame, candidatos: pd.DataFrame, bloqueos: pd.DataFrame, periodo: pd.DataFrame, nivel: pd.DataFrame, duracion: pd.DataFrame, escala: pd.DataFrame, nac: pd.DataFrame, output_paths: Dict[str, Path], desktop_dir: Path) -> str:
    lines = [
        "# Informe Exploracion Gobernanza Transversal MU / Avance",
        "",
        f"- Estado: {ESTADO}",
        f"- Declaracion carga: {DECLARACION_CARGA}",
        "- Fuentes originales modificadas: NO",
        "- Archivos de carga generados: NO",
        "- SIES_READY generado: NO",
        "- Reglas transversales aplicadas: NO",
        "",
        "## Dictamen Global",
        "",
        markdown_table(dictamen),
        "",
        "## Candidatos Transversales",
        "",
        markdown_table(candidatos[["Criterio", "Aplica a MU", "Aplica a Avance", "Aplica transversalmente", "Dictamen"]]),
        "",
        "## Criterios Bloqueados",
        "",
        markdown_table(bloqueos),
        "",
        "## Periodo Academico",
        "",
        markdown_table(periodo),
        "",
        "## Nivel / Duracion / Escala / Nacionalidad",
        "",
        "### Nivel",
        markdown_table(nivel.head(12)),
        "",
        "### Duracion",
        markdown_table(duracion.head(12)),
        "",
        "### Escala",
        markdown_table(escala.head(12)),
        "",
        "### Nacionalidad/Pais/Niveles",
        markdown_table(nac),
        "",
        "## Dictamen",
        "",
        "No existe habilitacion para reutilizar automaticamente reglas entre Matricula Unificada y Avance Curricular. Los criterios comparables quedan como candidatos o bloqueos hasta confirmacion documental/funcional.",
        "",
        "## Archivos",
        "",
        f"- Excel: {output_paths['excel']}",
        f"- Informe: {output_paths['informe']}",
        f"- Manifest: {output_paths['manifest']}",
        f"- Script: {output_paths['script']}",
        f"- Carpeta escritorio: {desktop_dir}",
        "",
    ]
    return "\n".join(lines)


def copy_to_desktop(output_dir: Path, desktop_dir: Path) -> None:
    desktop_dir.parent.mkdir(parents=True, exist_ok=True)
    if desktop_dir.exists():
        shutil.rmtree(desktop_dir)
    shutil.copytree(output_dir, desktop_dir, ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))


def print_console(dictamen: pd.DataFrame, candidatos: pd.DataFrame, bloqueos: pd.DataFrame, periodo: pd.DataFrame, nivel: pd.DataFrame, duracion: pd.DataFrame, escala: pd.DataFrame, nac: pd.DataFrame, output_paths: Dict[str, Path], desktop_dir: Path) -> None:
    print("EXPLORACIÓN GOBERNANZA TRANSVERSAL MU / AVANCE — GENERADA")
    print("Fuentes originales modificadas: NO")
    print("Declaración carga: NO_LISTO_PARA_CARGA")
    print("Estado: EXPLORATORIO")
    print()
    print("DICTAMEN GLOBAL")
    print(dictamen.to_string(index=False))
    print()
    print("CRITERIOS TRANSVERSALES CANDIDATOS")
    print(candidatos[["Criterio", "Aplica a MU", "Aplica a Avance", "Aplica transversalmente", "Dictamen"]].to_string(index=False))
    print()
    print("CRITERIOS BLOQUEADOS")
    print(bloqueos.to_string(index=False))
    print()
    print("PERIODO ACADÉMICO")
    print(periodo.to_string(index=False))
    print()
    print("NIVEL / DURACIÓN / ESCALA / NACIONALIDAD")
    print("NIVEL")
    print(nivel.head(10).to_string(index=False))
    print()
    print("DURACION")
    print(duracion.head(10).to_string(index=False))
    print()
    print("ESCALA")
    print(escala.head(10).to_string(index=False))
    print()
    print("NACIONALIDAD / PAIS / NIVELES")
    print(nac.to_string(index=False))
    print()
    print("ARCHIVOS")
    print(f"Excel: {output_paths['excel']}")
    print(f"Informe: {output_paths['informe']}")
    print(f"Manifest: {output_paths['manifest']}")
    print(f"Script: {output_paths['script']}")
    print(f"Carpeta escritorio: {desktop_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Exploracion gobernanza transversal MU / Avance")
    parser.add_argument("--timestamp", default=None, help="Timestamp YYYYMMDD_HHMMSS")
    args = parser.parse_args()
    timestamp = args.timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")

    output_dir = BASE_64 / f"EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    desktop_dir = Path.home() / "Desktop" / f"AVANCE_CURRICULAR_2026_EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE_{timestamp}"
    output_paths = {
        "excel": output_dir / OUTPUT_FILES["excel"],
        "informe": output_dir / OUTPUT_FILES["informe"],
        "manifest": output_dir / OUTPUT_FILES["manifest"],
        "script": output_dir / OUTPUT_FILES["script"],
    }
    current_script = Path(__file__).resolve()
    if current_script != output_paths["script"].resolve():
        shutil.copy2(current_script, output_paths["script"])

    archivos = build_archivos_explorados()
    manuales = extract_manual_rows()
    tsvs = build_tsv_inventory()
    candidatos = build_candidates()
    no_reuse = build_no_reuse()
    columnas = build_columnas_comparables()
    diccionarios = build_diccionarios_catalogos()
    periodo = build_periodo()
    nivel = build_nivel()
    duracion = build_duracion()
    escala = build_escala()
    nac = build_nac_pais()
    trace = build_control_trace()
    bloqueos = build_bloqueos()
    recomendacion = build_recomendacion()

    counts = {
        "archivos_explorados": len(archivos),
        "archivos_encontrados": int((archivos["Existe SI/NO"] == "SI").sum()),
        "archivos_faltantes": int((archivos["Existe SI/NO"] == "NO").sum()),
        "reglas_oficiales_detectadas": int((manuales["Regla textual o resumen trazable"] != "No se encontro evidencia local por busqueda de terminos.").sum()),
        "criterios_institucionales_detectados": len(candidatos),
        "criterios_tecnicos_detectados": int((archivos["Nivel de respaldo A/B/C/D/E"] == "C").sum()),
        "criterios_reutilizables_candidatos": int((candidatos["Aplica transversalmente"] == "PENDIENTE").sum()),
        "criterios_bloqueados": len(bloqueos),
    }
    dictamen = pd.DataFrame(
        [
            {
                "Proceso principal": PROCESO_PRINCIPAL,
                "Procesos comparados": PROCESOS_COMPARADOS,
                "Año proceso Matrícula Unificada": ANIO_MU,
                "Año proceso Avance Curricular": ANIO_AVANCE,
                "Estado": ESTADO,
                "Declaración carga": DECLARACION_CARGA,
                "Fuentes originales modificadas": "NO",
                "Archivos explorados": counts["archivos_explorados"],
                "Archivos encontrados": counts["archivos_encontrados"],
                "Archivos faltantes": counts["archivos_faltantes"],
                "Reglas oficiales detectadas": counts["reglas_oficiales_detectadas"],
                "Criterios institucionales detectados": counts["criterios_institucionales_detectados"],
                "Criterios técnicos detectados": counts["criterios_tecnicos_detectados"],
                "Criterios reutilizables candidatos": counts["criterios_reutilizables_candidatos"],
                "Criterios bloqueados": counts["criterios_bloqueados"],
                "Dictamen global": "Exploratorio: hay criterios comparables, pero ninguna regla transversal queda habilitada sin confirmacion documental/funcional.",
            }
        ]
    )

    sheets: Dict[str, pd.DataFrame] = OrderedDict(
        [
            ("00_DICTAMEN_GLOBAL", dictamen),
            ("01_ARCHIVOS_EXPLORADOS", archivos),
            ("02_MANUALES_OFICIALES", manuales),
            ("03_TSV_GOBERNANZA_LOCAL", tsvs),
            ("04_CRITERIOS_TRANSVERSALES_CANDIDATOS", candidatos),
            ("05_CRITERIOS_NO_REUTILIZAR", no_reuse),
            ("06_COLUMNAS_COMPARABLES_MU_AVANCE", columnas),
            ("07_DICCIONARIOS_Y_CATALOGOS", diccionarios),
            ("08_PERIODO_ACADEMICO", periodo),
            ("09_NIVEL", nivel),
            ("10_DURACION_ESTUDIOS", duracion),
            ("11_ESCALA_NOTAS", escala),
            ("12_NACIONALIDAD_PAIS_NIVELES", nac),
            ("13_CONTROL_TRACE", trace),
            ("14_BLOQUEOS_Y_PENDIENTES", bloqueos),
            ("15_RECOMENDACION_OPERATIVA", recomendacion),
        ]
    )
    # Fuentes y manifest se completan al final con hashes de salidas.
    placeholder_manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets, counts)
    sheets["16_FUENTES"] = build_fuentes(archivos, output_paths, desktop_dir)
    sheets["17_MANIFEST_LEGIBLE"] = manifest_legible(placeholder_manifest)
    write_excel(output_paths["excel"], sheets)

    output_paths["informe"].write_text(
        build_markdown(dictamen, candidatos, bloqueos, periodo, nivel, duracion, escala, nac, output_paths, desktop_dir),
        encoding="utf-8",
    )
    sheets["16_FUENTES"] = build_fuentes(archivos, output_paths, desktop_dir)
    manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets, counts)
    sheets["17_MANIFEST_LEGIBLE"] = manifest_legible(manifest)
    write_excel(output_paths["excel"], sheets)
    manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets, counts)
    output_paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    # Refrescar fuentes/manifest legible incluyendo manifest ya escrito.
    sheets["16_FUENTES"] = build_fuentes(archivos, output_paths, desktop_dir)
    sheets["17_MANIFEST_LEGIBLE"] = manifest_legible(manifest)
    write_excel(output_paths["excel"], sheets)
    output_paths["manifest"].write_text(json.dumps(build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets, counts), ensure_ascii=False, indent=2), encoding="utf-8")

    copy_to_desktop(output_dir, desktop_dir)
    print_console(dictamen, candidatos, bloqueos, periodo, nivel, duracion, escala, nac, output_paths, desktop_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
