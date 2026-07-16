#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Auditoria especifica de los 7 casos BLOQUEO_SIN_CODCLI_LISTA.

No modifica fuentes originales, no corrige datos, no genera SIES_READY,
no recalcula 20-21 y no mezcla hitos 54, 55 ni 56. Solo diagnostica causas y
propone rutas de resolucion.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Auditoria sin CODCLI_LISTA 5809"
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
NO_CORRECCION = "NO_APLICAR_CORRECCION_AUTOMATICA"
ESTADO_VALIDADOR_REQUERIDO = "APTO_PARA_CONTROL_PREVIO"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_57 = REPO / "avance_curricular_2026" / "57_auditoria_sin_codcli_lista_5809"

GOBERNANZA_51 = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)
VALIDADOR_52_SCRIPT = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "validador_gobernanza_precalculo_5809.py"
)
VALIDADOR_52_XLSX = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx"
)
MATRIZ_53 = (
    REPO
    / "avance_curricular_2026/53_matriz_resolucion_bloqueos_16_19_gobernada_5809/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_20260708_170521/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_5809.xlsx"
)
PRECARGA_5809 = (
    REPO
    / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/"
    / "5809_Precarga Matrícula Avance Curricular 2026.csv"
)
PROMEDIOS = (
    REPO
    / "avance_curricular_2026/25_actualizacion_fuente_promedios/"
    / "CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/00_FUENTE_CONGELADA/"
    / "PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
)
MAPEO_IDENTIDAD = (
    REPO
    / "avance_curricular_2026/05_cierre_integral/CIERRE_AVANCE_CURRICULAR_20260703_125157/"
    / "04_CONCILIACION_MATRICULA/PREPARACION_MATRICULA_20260703_141254/02_IDENTIFICADORES/"
    / "MAPEO_IDENTIDAD_CODCLI.xlsx"
)

OUTPUT_FILES = {
    "excel": "AUDITORIA_SIN_CODCLI_LISTA_5809.xlsx",
    "informe": "INFORME_AUDITORIA_SIN_CODCLI_LISTA_5809.md",
    "manifest": "manifest_auditoria_sin_codcli_lista_5809.json",
    "script": "auditoria_sin_codcli_lista_5809.py",
}

CAUSES = [
    "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD",
    "CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD",
    "RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO",
    "RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI",
    "DOCUMENTO_NO_CALZA_ENTRE_PRECARGA_Y_MAPEO",
    "POSIBLE_CASO_EXTRANJERO_O_DOCUMENTO_NO_RUT",
    "REQUIERE_MAPEO_INSTITUCIONAL",
    "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA",
    "REQUIERE_REVISION_FUNCIONAL",
]

SHEETS = OrderedDict(
    [
        ("00_DICTAMEN", "00_DICTAMEN"),
        ("01_VALIDACION_GOBERNANZA", "01_VALIDACION_GOBERNANZA"),
        ("02_RESUMEN_CAUSAS", "02_RESUMEN_CAUSAS"),
        ("03_7_CASOS", "03_7_CASOS"),
        ("04_BUSQUEDA_RUT_MAPEO", "04_BUSQUEDA_RUT_MAPEO"),
        ("05_BUSQUEDA_RUT_PROMEDIOS", "05_BUSQUEDA_RUT_PROMEDIOS"),
        ("06_BUSQUEDA_CODIGO_UNICO", "06_BUSQUEDA_CODIGO_UNICO"),
        ("07_ESTADO_ACADEMICO_OBSERVADO", "07_ESTADO_ACADEMICO_OBSERVADO"),
        ("08_CAUSAS_PROPUESTAS", "08_CAUSAS_PROPUESTAS"),
        ("09_NO_RESOLVER_AUTOMATICO", "09_NO_RESOLVER_AUTOMATICO"),
        ("10_FUENTES", "10_FUENTES"),
    ]
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def norm(value: Any) -> str:
    text = clean(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    out = []
    for ch in text.upper():
        out.append(ch if ch.isalnum() else "_")
    return "_".join("".join(out).split("_"))


def compact_number(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return text


def display(value: Any, max_len: int = 160) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")


def safe_read_excel(path: Path, sheet: str, columns: Iterable[str] | None = None) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")
    if columns is None:
        return df
    existing = [col for col in columns if col in df.columns]
    return df[existing].copy()


def to_int(value: Any) -> int:
    text = clean(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def normalize_body(value: Any) -> str:
    return "".join(ch for ch in clean(value).upper() if ch.isalnum())


def rut_body_from(value: Any) -> str:
    text = clean(value)
    if "-" in text:
        return normalize_body(text.split("-")[0])
    return normalize_body(text)


def join_unique(values: Iterable[Any], sep: str = " | ") -> str:
    return sep.join(sorted(set(clean(v) for v in values if clean(v))))


def split_pipe(value: Any) -> List[str]:
    text = clean(value)
    if not text:
        return []
    text = text.replace(";", "|")
    return [part.strip() for part in text.split("|") if part.strip()]


def require_paths() -> None:
    paths = [
        GOBERNANZA_51,
        VALIDADOR_52_SCRIPT,
        VALIDADOR_52_XLSX,
        MATRIZ_53,
        PRECARGA_5809,
        PROMEDIOS,
        MAPEO_IDENTIDAD,
    ]
    missing = [str(path) for path in paths if not path.exists()]
    if missing:
        raise FileNotFoundError("Faltan entradas obligatorias:\n" + "\n".join(missing))


def run_validator_52(output_dir: Path, timestamp: str) -> Dict[str, Any]:
    plan_path = output_dir / "plan_uso_validador_h52_auditoria_sin_codcli_lista.json"
    plan = {
        "descripcion": (
            "Auditoria diagnostica de 7 casos BLOQUEO_SIN_CODCLI_LISTA. "
            "No corrige, no recalcula 20-21, no genera carga y no mezcla hitos 54/55/56."
        ),
        "columnas_usadas": ["CODIGO_UNICO", "PLAN_ESTUDIOS", "NUM_DOCUMENTO", "DV", "CODCLI_LISTA", "CODCLI"],
        "llaves_usadas": ["CODIGO_UNICO", "PLAN_ESTUDIOS", "NUM_DOCUMENTO", "DV"],
        "diccionarios_usados": [],
        "transformaciones": [
            "Diagnostico documental de ausencia de CODCLI_LISTA; sin correccion ni remapeo automatico"
        ],
        "campos_calculo": [],
        "anio_calculo": 2025,
        "genera_sies_ready": False,
        "modifica_fuentes_originales": False,
        "usa_todos_codcli_rut": False,
        "filtra_por_programa": True,
        "usa_inferencias_como_reglas_oficiales": False,
        "usa_estados_AEIR": False,
        "usa_codcli_lista": True,
        "mezcla_anual_acumulado": False,
        "modo": "AUDITORIA_DIAGNOSTICA_NO_CORRECTIVA_SIN_CODCLI_LISTA",
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")

    validator_timestamp = f"{timestamp}_H57"
    cmd = [
        sys.executable,
        str(VALIDADOR_52_SCRIPT),
        "--timestamp",
        validator_timestamp,
        "--plan-json",
        str(plan_path),
    ]
    completed = subprocess.run(cmd, check=False, capture_output=True, text=True)
    validator_dir = (
        REPO
        / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809"
        / f"VALIDADOR_GOBERNANZA_PRECALCULO_5809_{validator_timestamp}"
    )
    manifest_path = validator_dir / "manifest_validacion_gobernanza_precalculo_5809.json"
    manifest: Dict[str, Any] = {}
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    result = {
        "plan_path": str(plan_path),
        "validator_timestamp": validator_timestamp,
        "validator_dir": str(validator_dir),
        "validator_manifest": str(manifest_path),
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "estado_validador": manifest.get("estado_validador", "SIN_MANIFEST"),
        "manifest": manifest,
    }
    if completed.returncode != 0:
        raise RuntimeError(
            "Validador hito 52 fallo.\nSTDOUT:\n"
            + completed.stdout
            + "\nSTDERR:\n"
            + completed.stderr
        )
    if result["estado_validador"] != ESTADO_VALIDADOR_REQUERIDO:
        raise RuntimeError(
            f"Validador hito 52 no apto: {result['estado_validador']}. "
            f"Manifest: {manifest_path}"
        )
    return result


def prepare_mapeo() -> pd.DataFrame:
    df = read_sheet(MAPEO_IDENTIDAD, "MAPEO")
    df["ID_FILA_5809_N"] = df.get("ID_FILA_5809", "").map(clean)
    df["CODIGO_UNICO_N"] = df.get("CODIGO_UNICO", "").map(clean)
    df["DOC_PRECARGA_BODY_N"] = df.get("DOC_PRECARGA_BODY", "").map(normalize_body)
    df["NUM_DOCUMENTO_N"] = df.get("NUM_DOCUMENTO", "").map(normalize_body)
    df["RUT_INST_BODY_N"] = df.get("RUT_INST_BODY", "").map(normalize_body)
    df["CODCLI_LISTA_N"] = df.get("CODCLI_LISTA", "").map(clean)
    return df


def prepare_precarga() -> pd.DataFrame:
    try:
        df = pd.read_csv(PRECARGA_5809, sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    except UnicodeDecodeError:
        df = pd.read_csv(PRECARGA_5809, sep=";", dtype=str, encoding="latin1").fillna("")
    df.insert(0, "ID_FILA_5809", [str(i) for i in range(1, len(df) + 1)])
    df["ID_FILA_5809_N"] = df["ID_FILA_5809"].map(clean)
    df["NUM_DOCUMENTO_N"] = df.get("NUM_DOCUMENTO", "").map(normalize_body)
    df["CODIGO_UNICO_N"] = df.get("CODIGO_UNICO", "").map(clean)
    return df


def prepare_promedios() -> Dict[str, pd.DataFrame]:
    hoja1_cols = [
        "CODCLI",
        "RUT",
        "DIG",
        "ANO",
        "PERIODO",
        "ESTADO",
        "DESCRIPCION_ESTADO",
        "CODRAMO",
        "ESTADO_ACADEMICO",
        "CODCARR",
        "NIVEL",
    ]
    alumnos_cols = [
        "CODCLI",
        "CODCARPR",
        "ANOMATRICULA",
        "PERIODOMATRICULA",
        "RUT",
        "ESTADOACADEMICO",
        "SITUACION",
        "MATRICULA",
        "NIVEL",
    ]
    base_cols = ["RUT", "DV", "CODCLI"]
    out = {
        "Hoja1": safe_read_excel(PROMEDIOS, "Hoja1", hoja1_cols),
        "DatosAlumnos": safe_read_excel(PROMEDIOS, "DatosAlumnos", alumnos_cols),
        "Matricula_2025": safe_read_excel(PROMEDIOS, "Matricula_2025", alumnos_cols),
        "base_datos": safe_read_excel(PROMEDIOS, "base_datos", base_cols),
    }
    for name, df in out.items():
        if "CODCLI" in df.columns:
            df["CODCLI_N"] = df["CODCLI"].map(clean)
        else:
            df["CODCLI_N"] = ""
        if "RUT" in df.columns:
            df["RUT_BODY_N"] = df["RUT"].map(rut_body_from)
        else:
            df["RUT_BODY_N"] = ""
        if "ANO" in df.columns:
            df["ANO_N"] = df["ANO"].map(compact_number)
        if "ANOMATRICULA" in df.columns:
            df["ANOMATRICULA_N"] = df["ANOMATRICULA"].map(compact_number)
        out[name] = df
    return out


def source_summary(source: str, df: pd.DataFrame, body: str) -> Dict[str, Any]:
    sub = df[df["RUT_BODY_N"] == body].copy() if body else df.iloc[0:0].copy()
    years: List[Any] = []
    if "ANO_N" in sub.columns:
        years.extend(sub["ANO_N"].tolist())
    if "ANOMATRICULA_N" in sub.columns:
        years.extend(sub["ANOMATRICULA_N"].tolist())
    estados: List[Any] = []
    for col in ["ESTADO_ACADEMICO", "ESTADOACADEMICO"]:
        if col in sub.columns:
            estados.extend(sub[col].tolist())
    codcarr: List[Any] = []
    for col in ["CODCARR", "CODCARPR"]:
        if col in sub.columns:
            codcarr.extend(sub[col].tolist())
    return {
        "Fuente": source,
        "Filas fuente": len(sub),
        "CODCLI observados": join_unique(sub["CODCLI_N"]) if "CODCLI_N" in sub.columns else "",
        "CODCARR/CODCARPR": join_unique(codcarr),
        "Anos observados": join_unique(years),
        "Filas 2025": int((sub["ANO_N"] == "2025").sum()) if "ANO_N" in sub.columns else 0,
        "Estados academicos": join_unique(estados),
        "Situaciones": join_unique(sub["SITUACION"]) if "SITUACION" in sub.columns else "",
        "Estados ramo": join_unique(sub["ESTADO"]) if "ESTADO" in sub.columns else "",
        "Descripciones estado": join_unique(sub["DESCRIPCION_ESTADO"]) if "DESCRIPCION_ESTADO" in sub.columns else "",
        "CODRAMOS distintos": sub["CODRAMO"].nunique() if "CODRAMO" in sub.columns and not sub.empty else 0,
    }


def route_for_cause(cause: str) -> str:
    routes = {
        "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD": "Solicitar mapeo institucional de identidad antes de cualquier conteo.",
        "CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD": "Validar CODIGO_UNICO contra precarga/carreras y fuente institucional.",
        "RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO": "Revisar programa/carrera antes de asociar CODCLI; no sumar por RUT.",
        "RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI": "Revisar identidad y programa; no usar el CODCLI encontrado sin mapeo funcional.",
        "DOCUMENTO_NO_CALZA_ENTRE_PRECARGA_Y_MAPEO": "Resolver contradiccion documental entre precarga y mapeo institucional.",
        "POSIBLE_CASO_EXTRANJERO_O_DOCUMENTO_NO_RUT": "Validar documento no RUT o extranjero con registro institucional.",
        "REQUIERE_MAPEO_INSTITUCIONAL": "Completar CODCLI_LISTA gobernado o documentar decision funcional de ausencia.",
        "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA": "Solicitar fuente academica complementaria si no existe CODCLI trazable.",
        "REQUIERE_REVISION_FUNCIONAL": "Resolver con responsable funcional antes de cualquier correccion.",
    }
    return routes[cause]


def classify_case(
    row: pd.Series,
    mapeo: pd.DataFrame,
    precarga: pd.DataFrame,
    promedios: Dict[str, pd.DataFrame],
) -> Tuple[Dict[str, Any], Dict[str, Any], Dict[str, Any], Dict[str, Any], List[Dict[str, Any]], Dict[str, Any]]:
    fila = clean(row.get("FILA_5809", ""))
    body = rut_body_from(row.get("RUT_NORMALIZADO", ""))
    codigo = clean(row.get("CODIGO_UNICO", ""))
    mcase = mapeo[mapeo["ID_FILA_5809_N"] == fila].copy()
    pcase = precarga[precarga["ID_FILA_5809_N"] == fila].copy()

    same_doc = mapeo[
        (mapeo["DOC_PRECARGA_BODY_N"] == body)
        | (mapeo["NUM_DOCUMENTO_N"] == body)
        | (mapeo["RUT_INST_BODY_N"] == body)
    ].copy()
    identity_doc = same_doc[(same_doc["RUT_INST_BODY_N"] == body) & (same_doc["RUT_INST_BODY_N"] != "")]
    same_code = mapeo[mapeo["CODIGO_UNICO_N"] == codigo].copy()
    code_identity = same_code[(same_code["RUT_INST_BODY_N"] != "") | (same_code["CODCLI_LISTA_N"] != "")]
    identity_same_code = identity_doc[identity_doc["CODIGO_UNICO_N"] == codigo]

    source_rows = []
    prom_codcli_values: List[str] = []
    for source, df in promedios.items():
        summary = source_summary(source, df, body)
        summary["FILA_5809"] = fila
        summary["RUT_BUSCADO"] = clean(row.get("RUT_NORMALIZADO", ""))
        source_rows.append(summary)
        prom_codcli_values.extend(split_pipe(summary["CODCLI observados"]))
    prom_codcli_values = [value for value in prom_codcli_values if value]

    mrow = mcase.iloc[0] if not mcase.empty else pd.Series(dtype=object)
    tipo_doc = clean(mrow.get("TIPO_DOCUMENTO", ""))
    doc_mismatch = (
        "CONTRADICCION_DOCUMENTO" in norm(mrow.get("CLASIFICACION_IDENTIDAD", ""))
        or norm(mrow.get("COINCIDE_DV_INSTITUCIONAL", "")) == "NO"
    )
    possible_foreign = bool(tipo_doc and norm(tipo_doc) != "R") or not clean(mrow.get("DV", ""))
    rut_no_identity = identity_doc.empty
    code_no_exists = same_code.empty
    rut_exists_not_code = (not identity_doc.empty) and identity_same_code.empty
    prom_other_codcli = bool(prom_codcli_values)
    requires_complement = (not prom_other_codcli) or doc_mismatch

    flags = {
        "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD": bool(rut_no_identity),
        "CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD": bool(code_no_exists),
        "RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO": bool(rut_exists_not_code),
        "RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI": bool(prom_other_codcli),
        "DOCUMENTO_NO_CALZA_ENTRE_PRECARGA_Y_MAPEO": bool(doc_mismatch),
        "POSIBLE_CASO_EXTRANJERO_O_DOCUMENTO_NO_RUT": bool(possible_foreign),
        "REQUIERE_MAPEO_INSTITUCIONAL": True,
        "REQUIERE_FUENTE_ACADEMICA_COMPLEMENTARIA": bool(requires_complement),
        "REQUIERE_REVISION_FUNCIONAL": True,
    }

    if flags["DOCUMENTO_NO_CALZA_ENTRE_PRECARGA_Y_MAPEO"]:
        principal = "DOCUMENTO_NO_CALZA_ENTRE_PRECARGA_Y_MAPEO"
    elif flags["RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI"]:
        principal = "RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI"
    elif flags["POSIBLE_CASO_EXTRANJERO_O_DOCUMENTO_NO_RUT"]:
        principal = "POSIBLE_CASO_EXTRANJERO_O_DOCUMENTO_NO_RUT"
    elif flags["RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD"]:
        principal = "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD"
    elif flags["CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD"]:
        principal = "CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD"
    elif flags["RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO"]:
        principal = "RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO"
    else:
        principal = "REQUIERE_MAPEO_INSTITUCIONAL"

    detected = [cause for cause in CAUSES if flags[cause]]
    case_row = {
        "FILA_5809": fila,
        "RUT_NORMALIZADO": clean(row.get("RUT_NORMALIZADO", "")),
        "CODIGO_UNICO": codigo,
        "PLAN_ESTUDIOS": clean(row.get("PLAN_ESTUDIOS", "")),
        "TIPO_DOCUMENTO_OBSERVADO": tipo_doc,
        "NUM_DOCUMENTO_MAPEO": clean(mrow.get("NUM_DOCUMENTO", "")),
        "DV_MAPEO": clean(mrow.get("DV", "")),
        "CODCLI_LISTA_NORM": clean(row.get("CODCLI_LISTA_NORM", "")),
        "DICTAMEN_HITO_53": clean(row.get("DICTAMEN_DICCIONARIO_ESTADO", "")),
        "CLASIFICACION_MAPEO": clean(mrow.get("CLASIFICACION_IDENTIDAD", "")),
        "CAUSA_PRINCIPAL_PROPUESTA": principal,
        "CAUSAS_DETECTADAS": " | ".join(detected),
        "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD": "SI" if rut_no_identity else "NO",
        "RUT_PROMEDIOS_CODCLI_OBSERVADO": join_unique(prom_codcli_values),
        "DOCUMENTO_NO_CALZA": "SI" if doc_mismatch else "NO",
        "POSIBLE_DOCUMENTO_NO_RUT": "SI" if possible_foreign else "NO",
        "ACCION_PROPUESTA": route_for_cause(principal),
        "PUEDE_RESOLVERSE_AUTOMATICAMENTE_AHORA": "NO",
        "NO_APLICAR_CORRECCION": NO_CORRECCION,
        "DECLARACION_CARGA": DECLARACION_CARGA,
    }

    mapeo_row = {
        "FILA_5809": fila,
        "RUT_BUSCADO": clean(row.get("RUT_NORMALIZADO", "")),
        "DOC_PRECARGA_BODY_MAPEO": clean(mrow.get("DOC_PRECARGA_BODY", "")),
        "RUT_INST_BODY_MAPEO": clean(mrow.get("RUT_INST_BODY", "")),
        "N_REGISTROS_DATOS": clean(mrow.get("N_REGISTROS_DATOS", "")),
        "N_CODCLI_OBSERVADO_NO_USAR_COMO_LLAVE": clean(mrow.get("N_CODCLI", "")),
        "CODCLI_LISTA_MAPEO": clean(mrow.get("CODCLI_LISTA", "")),
        "CODCARPR_LISTA_MAPEO": clean(mrow.get("CODCARPR_LISTA", "")),
        "COINCIDE_DV_INSTITUCIONAL": clean(mrow.get("COINCIDE_DV_INSTITUCIONAL", "")),
        "CODCARR_ESPERADO_EN_DATOS": clean(mrow.get("CODCARR_ESPERADO_EN_DATOS", "")),
        "CLASIFICACION_IDENTIDAD": clean(mrow.get("CLASIFICACION_IDENTIDAD", "")),
        "FILAS_MAPEO_MISMO_DOCUMENTO": len(same_doc),
        "FILAS_MAPEO_IDENTIDAD_DOCUMENTO": len(identity_doc),
        "RUT_NO_EXISTE_EN_MAPEO_IDENTIDAD": "SI" if rut_no_identity else "NO",
        "USO": "Evidencia diagnostica; no corrige ni crea CODCLI_LISTA.",
    }

    prom_row = {
        "FILA_5809": fila,
        "RUT_BUSCADO": clean(row.get("RUT_NORMALIZADO", "")),
        "CODCLI_PROMEDIOS_OBSERVADOS": join_unique(prom_codcli_values),
        "RUT_EXISTE_EN_PROMEDIOS_CON_OTRO_CODCLI": "SI" if prom_other_codcli else "NO",
        "HOJA1_FILAS": next(r["Filas fuente"] for r in source_rows if r["Fuente"] == "Hoja1"),
        "HOJA1_FILAS_2025": next(r["Filas 2025"] for r in source_rows if r["Fuente"] == "Hoja1"),
        "DATOSALUMNOS_FILAS": next(r["Filas fuente"] for r in source_rows if r["Fuente"] == "DatosAlumnos"),
        "MATRICULA_2025_FILAS": next(r["Filas fuente"] for r in source_rows if r["Fuente"] == "Matricula_2025"),
        "BASE_DATOS_FILAS": next(r["Filas fuente"] for r in source_rows if r["Fuente"] == "base_datos"),
        "USO": "Evidencia diagnostica; no habilita correccion sin mapeo CODCLI_LISTA.",
    }

    codigo_row = {
        "FILA_5809": fila,
        "CODIGO_UNICO": codigo,
        "FILAS_MAPEO_CODIGO_UNICO": len(same_code),
        "FILAS_MAPEO_CODIGO_UNICO_CON_IDENTIDAD": len(code_identity),
        "FILAS_MAPEO_CODIGO_UNICO_MISMO_DOCUMENTO": len(identity_same_code),
        "CODIGO_UNICO_NO_EXISTE_EN_MAPEO_IDENTIDAD": "SI" if code_no_exists else "NO",
        "RUT_EXISTE_PERO_NO_PARA_CODIGO_UNICO": "SI" if rut_exists_not_code else "NO",
        "DOCUMENTOS_CODIGO_UNICO": join_unique(same_code["DOC_PRECARGA_BODY_N"]) if not same_code.empty else "",
        "CLASIFICACIONES_CODIGO_UNICO": join_unique(same_code["CLASIFICACION_IDENTIDAD"]) if not same_code.empty else "",
        "USO": "Evidencia de programa; no resuelve identidad por nombre ni por codigo solo.",
    }

    no_auto_row = {
        "FILA_5809": fila,
        "RUT_NORMALIZADO": clean(row.get("RUT_NORMALIZADO", "")),
        "CODIGO_UNICO": codigo,
        "CAUSA_PRINCIPAL_PROPUESTA": principal,
        "MOTIVO_NO_AUTOMATICO": (
            "No existe CODCLI_LISTA gobernado para la fila; cualquier asociacion requiere mapeo institucional "
            "o fuente academica complementaria validada funcionalmente."
        ),
        "NO_APLICAR_CORRECCION": NO_CORRECCION,
        "DECLARACION_CARGA": DECLARACION_CARGA,
    }
    return case_row, mapeo_row, prom_row, codigo_row, source_rows, no_auto_row


def build_case_tables(
    cases: pd.DataFrame,
    mapeo: pd.DataFrame,
    precarga: pd.DataFrame,
    promedios: Dict[str, pd.DataFrame],
) -> Dict[str, pd.DataFrame]:
    case_rows: List[Dict[str, Any]] = []
    mapeo_rows: List[Dict[str, Any]] = []
    prom_rows: List[Dict[str, Any]] = []
    codigo_rows: List[Dict[str, Any]] = []
    estado_rows: List[Dict[str, Any]] = []
    no_auto_rows: List[Dict[str, Any]] = []
    for _, row in cases.iterrows():
        case_row, mapeo_row, prom_row, codigo_row, source_rows, no_auto_row = classify_case(
            row, mapeo, precarga, promedios
        )
        case_rows.append(case_row)
        mapeo_rows.append(mapeo_row)
        prom_rows.append(prom_row)
        codigo_rows.append(codigo_row)
        for item in source_rows:
            item["RUT_NORMALIZADO"] = clean(row.get("RUT_NORMALIZADO", ""))
            item["CODIGO_UNICO"] = clean(row.get("CODIGO_UNICO", ""))
            estado_rows.append(item)
        no_auto_rows.append(no_auto_row)
    return {
        "casos": pd.DataFrame(case_rows),
        "rut_mapeo": pd.DataFrame(mapeo_rows),
        "rut_promedios": pd.DataFrame(prom_rows),
        "codigo_unico": pd.DataFrame(codigo_rows),
        "estado_observado": pd.DataFrame(estado_rows),
        "no_auto": pd.DataFrame(no_auto_rows),
    }


def build_resumen_causas(cases_df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for cause in CAUSES:
        rows.append(
            {
                "CAUSA": cause,
                "CASOS_CAUSA_PRINCIPAL": int((cases_df["CAUSA_PRINCIPAL_PROPUESTA"] == cause).sum()),
                "CASOS_CAUSA_DETECTADA": int(cases_df["CAUSAS_DETECTADAS"].str.contains(cause, regex=False).sum()),
                "RESOLVER_AUTOMATICAMENTE": "NO",
                "RUTA_PROPUESTA": route_for_cause(cause),
                "DECLARACION_CARGA": DECLARACION_CARGA,
            }
        )
    return pd.DataFrame(rows)


def build_causas_propuestas(resumen: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "CAUSA": row["CAUSA"],
                "CASOS_CAUSA_PRINCIPAL": row["CASOS_CAUSA_PRINCIPAL"],
                "CASOS_CAUSA_DETECTADA": row["CASOS_CAUSA_DETECTADA"],
                "DECISION_PROPUESTA": route_for_cause(row["CAUSA"]),
                "PUEDE_RESOLVERSE_AUTOMATICAMENTE_AHORA": "NO",
                "REQUIERE_REVISION_FUNCIONAL": "SI",
                "NO_APLICAR_CORRECCION": NO_CORRECCION,
            }
            for _, row in resumen.iterrows()
        ]
    )


def build_validation_df(validator: Dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Control": "Gobernanza integral hito 51",
                "Estado": "USADA",
                "Evidencia": str(GOBERNANZA_51),
                "Hash SHA256": sha256_file(GOBERNANZA_51),
                "Observacion": "Base obligatoria; CODCLI_LISTA gobernada, N_CODCLI bloqueada.",
            },
            {
                "Control": "Validador hito 52",
                "Estado": validator["estado_validador"],
                "Evidencia": validator["validator_manifest"],
                "Hash SHA256": sha256_file(Path(validator["validator_manifest"]))
                if Path(validator["validator_manifest"]).exists()
                else "",
                "Observacion": "Compuerta ejecutada antes de generar auditoria hito 57.",
            },
            {
                "Control": "Matriz hito 53",
                "Estado": "USADA",
                "Evidencia": str(MATRIZ_53),
                "Hash SHA256": sha256_file(MATRIZ_53),
                "Observacion": "Se usa solo hoja 06_SIN_CODCLI_LISTA; no se mezclan hitos 54/55/56.",
            },
            {
                "Control": "Precarga oficial 5809",
                "Estado": "USADA_COMO_EVIDENCIA",
                "Evidencia": str(PRECARGA_5809),
                "Hash SHA256": sha256_file(PRECARGA_5809),
                "Observacion": "Comparacion documental contra MAPEO; no se modifica.",
            },
            {
                "Control": "MAPEO_IDENTIDAD_CODCLI",
                "Estado": "USADA",
                "Evidencia": str(MAPEO_IDENTIDAD),
                "Hash SHA256": sha256_file(MAPEO_IDENTIDAD),
                "Observacion": "Busca CODCLI_LISTA y clasificacion de identidad; N_CODCLI no se usa como llave.",
            },
            {
                "Control": "PROMEDIOS actualizado",
                "Estado": "USADO_COMO_EVIDENCIA",
                "Evidencia": str(PROMEDIOS),
                "Hash SHA256": sha256_file(PROMEDIOS),
                "Observacion": "Busqueda diagnostica por documento/RUT; no habilita correccion automatica.",
            },
            {
                "Control": "Plan hito 52",
                "Estado": validator["estado_validador"],
                "Evidencia": validator["plan_path"],
                "Hash SHA256": sha256_file(Path(validator["plan_path"])),
                "Observacion": "Plan no correctivo, sin SIES_READY y sin modificar fuentes.",
            },
        ]
    )


def build_fuentes_df(output_dir: Path, output_paths: Dict[str, Path], validator: Dict[str, Any]) -> pd.DataFrame:
    inputs = [
        ("INPUT", "Gobernanza integral hito 51", GOBERNANZA_51),
        ("INPUT", "Validador hito 52 xlsx base", VALIDADOR_52_XLSX),
        ("INPUT", "Validador hito 52 script", VALIDADOR_52_SCRIPT),
        ("INPUT", "Matriz hito 53", MATRIZ_53),
        ("INPUT", "Precarga oficial 5809", PRECARGA_5809),
        ("INPUT", "PROMEDIOS actualizado", PROMEDIOS),
        ("INPUT", "MAPEO_IDENTIDAD_CODCLI", MAPEO_IDENTIDAD),
        ("INPUT_DERIVADO", "Manifest validador hito 52 ejecucion hito 57", Path(validator["validator_manifest"])),
    ]
    rows = []
    for tipo, nombre, path in inputs:
        rows.append(
            {
                "Tipo": tipo,
                "Nombre": nombre,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() else "",
                "Uso": "Entrada o compuerta obligatoria; no modificada.",
            }
        )
    for key, path in output_paths.items():
        rows.append(
            {
                "Tipo": "OUTPUT",
                "Nombre": key,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() and key != "manifest" else "",
                "Uso": "Producto hito 57 no apto para carga. Manifest sin hash autorreferencial.",
            }
        )
    rows.append(
        {
            "Tipo": "OUTPUT_DIR",
            "Nombre": "Carpeta hito 57",
            "Ruta": str(output_dir),
            "Existe": "SI" if output_dir.exists() else "NO",
            "Hash SHA256": "",
            "Uso": "Contenedor de auditoria especifica.",
        }
    )
    return pd.DataFrame(rows)


def write_excel(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for logical_name, df in sheets.items():
            df.to_excel(writer, sheet_name=SHEETS.get(logical_name, logical_name)[:31], index=False)
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            values = [display(cell.value, 90) for cell in column_cells[:200]]
            width = min(max([len(v) for v in values] + [10]) + 2, 60)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
    wb.save(path)


def build_markdown(
    output_dir: Path,
    dictamen_df: pd.DataFrame,
    validacion_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    casos_df: pd.DataFrame,
    manifest: Dict[str, Any],
) -> str:
    principal = resumen_df[resumen_df["CASOS_CAUSA_PRINCIPAL"] > 0][
        ["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RUTA_PROPUESTA"]
    ]
    lines = [
        "# Informe Auditoria BLOQUEO_SIN_CODCLI_LISTA 5809",
        "",
        f"- Proceso: {PROCESO}",
        f"- Subproyecto: {SUBPROYECTO}",
        f"- Anio proceso: {ANIO_PROCESO}",
        f"- Anio referencia datos: {ANIO_REFERENCIA_DATOS}",
        f"- Declaracion carga: {DECLARACION_CARGA}",
        "- Fuentes originales modificadas: NO",
        "- Correcciones aplicadas: NO",
        "- SIES_READY generado: NO",
        "- Recalculo 20-21: NO",
        "- Mezcla con hitos 54/55/56: NO",
        f"- Carpeta: {output_dir}",
        "",
        "## Dictamen",
        "",
        markdown_table(dictamen_df),
        "",
        "## Validacion de gobernanza",
        "",
        markdown_table(validacion_df[["Control", "Estado", "Observacion"]]),
        "",
        "## Resumen de causas",
        "",
        markdown_table(resumen_df[["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RESOLVER_AUTOMATICAMENTE"]]),
        "",
        "## Causas principales con ruta",
        "",
        markdown_table(principal),
        "",
        "## Hallazgos",
        "",
        "- Los 7 casos provienen exclusivamente de la hoja 06_SIN_CODCLI_LISTA del hito 53.",
        "- No se usa N_CODCLI como CODCLI academico.",
        "- RUT y TIPO_DOCUMENTO se muestran solo como evidencia diagnostica; no son llave de calculo en esta auditoria.",
        "- Ningun caso queda habilitado para correccion automatica.",
        "",
        "## Casos auditados",
        "",
        markdown_table(
            casos_df[
                [
                    "FILA_5809",
                    "RUT_NORMALIZADO",
                    "CODIGO_UNICO",
                    "CLASIFICACION_MAPEO",
                    "CAUSA_PRINCIPAL_PROPUESTA",
                    "RUT_PROMEDIOS_CODCLI_OBSERVADO",
                ]
            ]
        ),
        "",
        "## Archivos",
        "",
        f"- Excel: {manifest.get('archivos', {}).get('excel', '')}",
        f"- Informe: {manifest.get('archivos', {}).get('informe', '')}",
        f"- Manifest: {manifest.get('archivos', {}).get('manifest', '')}",
        f"- Script: {manifest.get('archivos', {}).get('script', '')}",
        "",
    ]
    return "\n".join(lines)


def build_manifest(
    timestamp: str,
    output_dir: Path,
    output_paths: Dict[str, Path],
    validator: Dict[str, Any],
    dictamen_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    sheets: Dict[str, pd.DataFrame],
) -> Dict[str, Any]:
    return {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": timestamp,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "correcciones_aplicadas": "NO",
        "sies_ready_generado": "NO",
        "recalculo_20_21": "NO",
        "mezcla_hitos_54_55_56": "NO",
        "estado_validador_hito_52": validator["estado_validador"],
        "carpeta_salida": str(output_dir),
        "archivos": {key: str(path) for key, path in output_paths.items()},
        "entradas_obligatorias": {
            "gobernanza_51": str(GOBERNANZA_51),
            "validador_52_script": str(VALIDADOR_52_SCRIPT),
            "validador_52_xlsx": str(VALIDADOR_52_XLSX),
            "matriz_53": str(MATRIZ_53),
            "precarga_5809": str(PRECARGA_5809),
            "promedios_actualizado": str(PROMEDIOS),
            "mapeo_identidad_codcli": str(MAPEO_IDENTIDAD),
        },
        "validador_hito_52_ejecucion": {
            "plan": validator["plan_path"],
            "manifest": validator["validator_manifest"],
            "carpeta": validator["validator_dir"],
            "returncode": validator["returncode"],
        },
        "conteos": {
            "casos_auditados": int(dictamen_df.loc[0, "Total casos auditados"]),
            "casos_hoja_06_hito_53": int(dictamen_df.loc[0, "Casos fuente hito 53"]),
            "causas_principales": resumen_df.set_index("CAUSA")["CASOS_CAUSA_PRINCIPAL"].astype(int).to_dict(),
            "causas_detectadas": resumen_df.set_index("CAUSA")["CASOS_CAUSA_DETECTADA"].astype(int).to_dict(),
            "hojas_excel": {name: len(df) for name, df in sheets.items()},
        },
        "hashes_salida": {
            key: sha256_file(path) for key, path in output_paths.items() if path.exists() and key != "manifest"
        },
    }


def print_console(
    dictamen_df: pd.DataFrame,
    resumen_df: pd.DataFrame,
    output_paths: Dict[str, Path],
    output_dir: Path,
    validator_state: str,
) -> None:
    print("AUDITORIA SIN CODCLI_LISTA 5809 — GENERADA")
    print("Fuentes originales modificadas: NO")
    print("Declaracion carga: NO_LISTO_PARA_CARGA")
    print(f"Estado validador: {validator_state}")
    print()
    print("DICTAMEN")
    print(dictamen_df.to_string(index=False))
    print()
    print("RESUMEN CAUSAS")
    print(
        resumen_df[
            ["CAUSA", "CASOS_CAUSA_PRINCIPAL", "CASOS_CAUSA_DETECTADA", "RESOLVER_AUTOMATICAMENTE"]
        ].to_string(index=False)
    )
    print()
    print("ARCHIVOS")
    print(f"Excel: {output_paths['excel']}")
    print(f"Informe: {output_paths['informe']}")
    print(f"Manifest: {output_paths['manifest']}")
    print(f"Script: {output_paths['script']}")
    print(f"Carpeta hito: {output_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Auditoria hito 57 sin CODCLI_LISTA 5809")
    parser.add_argument("--timestamp", default=None, help="Timestamp YYYYMMDD_HHMMSS. Si se omite, usa fecha/hora actual.")
    args = parser.parse_args()

    timestamp = args.timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = BASE_57 / f"AUDITORIA_SIN_CODCLI_LISTA_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_paths = {
        "excel": output_dir / OUTPUT_FILES["excel"],
        "informe": output_dir / OUTPUT_FILES["informe"],
        "manifest": output_dir / OUTPUT_FILES["manifest"],
        "script": output_dir / OUTPUT_FILES["script"],
    }

    require_paths()
    current_script = Path(__file__).resolve()
    target_script = output_paths["script"].resolve()
    if current_script != target_script:
        shutil.copy2(current_script, target_script)

    validator = run_validator_52(output_dir, timestamp)

    matriz_06 = read_sheet(MATRIZ_53, "06_SIN_CODCLI_LISTA")
    cases = matriz_06[matriz_06["DICTAMEN_DICCIONARIO_ESTADO"] == "BLOQUEO_SIN_CODCLI_LISTA"].copy()
    if len(cases) != 7:
        raise RuntimeError(f"Se esperaban 7 casos BLOQUEO_SIN_CODCLI_LISTA; encontrados: {len(cases)}")

    mapeo = prepare_mapeo()
    precarga = prepare_precarga()
    promedios = prepare_promedios()
    case_tables = build_case_tables(cases, mapeo, precarga, promedios)
    casos_df = case_tables["casos"]
    resumen_df = build_resumen_causas(casos_df)
    causas_df = build_causas_propuestas(resumen_df)
    validacion_df = build_validation_df(validator)

    dictamen_df = pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Estado de carga": DECLARACION_CARGA,
                "Estado validador": validator["estado_validador"],
                "Fuentes originales modificadas": "NO",
                "Correcciones aplicadas": "NO",
                "SIES_READY generado": "NO",
                "Recalculo 20-21": "NO",
                "Mezcla hitos 54/55/56": "NO",
                "Casos fuente hito 53": len(cases),
                "Total casos auditados": len(casos_df),
                "Casos sin CODCLI_LISTA": int((casos_df["CODCLI_LISTA_NORM"].map(clean) == "").sum()),
                "Casos con posible documento no RUT": int((casos_df["POSIBLE_DOCUMENTO_NO_RUT"] == "SI").sum()),
                "Casos con documento no calza": int((casos_df["DOCUMENTO_NO_CALZA"] == "SI").sum()),
                "Casos con RUT en PROMEDIOS otro CODCLI": int(
                    (casos_df["RUT_PROMEDIOS_CODCLI_OBSERVADO"].map(clean) != "").sum()
                ),
                "Dictamen global": (
                    "NO_LISTO_PARA_CARGA: auditoria diagnostica de 7 casos sin CODCLI_LISTA; "
                    "no procede correccion automatica ni carga."
                ),
                "Proximo paso permitido": (
                    "Resolver mapeo institucional/documental o fuente academica complementaria; luego revalidar gobernanza."
                ),
            }
        ]
    )

    sheets = OrderedDict(
        [
            ("00_DICTAMEN", dictamen_df),
            ("01_VALIDACION_GOBERNANZA", validacion_df),
            ("02_RESUMEN_CAUSAS", resumen_df),
            ("03_7_CASOS", casos_df),
            ("04_BUSQUEDA_RUT_MAPEO", case_tables["rut_mapeo"]),
            ("05_BUSQUEDA_RUT_PROMEDIOS", case_tables["rut_promedios"]),
            ("06_BUSQUEDA_CODIGO_UNICO", case_tables["codigo_unico"]),
            ("07_ESTADO_ACADEMICO_OBSERVADO", case_tables["estado_observado"]),
            ("08_CAUSAS_PROPUESTAS", causas_df),
            ("09_NO_RESOLVER_AUTOMATICO", case_tables["no_auto"]),
        ]
    )
    sheets["10_FUENTES"] = build_fuentes_df(output_dir, output_paths, validator)
    write_excel(output_paths["excel"], sheets)

    manifest = build_manifest(timestamp, output_dir, output_paths, validator, dictamen_df, resumen_df, sheets)
    output_paths["informe"].write_text(
        build_markdown(output_dir, dictamen_df, validacion_df, resumen_df, casos_df, manifest),
        encoding="utf-8",
    )
    sheets["10_FUENTES"] = build_fuentes_df(output_dir, output_paths, validator)
    write_excel(output_paths["excel"], sheets)
    manifest = build_manifest(timestamp, output_dir, output_paths, validator, dictamen_df, resumen_df, sheets)
    output_paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print_console(dictamen_df, resumen_df, output_paths, output_dir, validator["estado_validador"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
