#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Validador previo obligatorio de gobernanza para Avance Curricular 5809.

Este script lee la gobernanza integral del paso 51 y valida que una futura
ejecucion de calculo/correccion no use fuentes, hojas, columnas, llaves,
diccionarios o transformaciones no gobernadas. No modifica fuentes originales,
no genera SIES_READY y no recalcula datos finales.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Validador gobernanza precalculo 5809"
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
ESTADO_APTO = "APTO_PARA_CONTROL_PREVIO"
ESTADO_BLOQUEADO = "BLOQUEADO"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_52 = REPO / "avance_curricular_2026" / "52_validador_gobernanza_precalculo_5809"
GOBERNANZA_XLSX = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)

OUTPUT_FILES = {
    "excel": "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx",
    "informe": "INFORME_VALIDACION_GOBERNANZA_PRECALCULO_5809.md",
    "manifest": "manifest_validacion_gobernanza_precalculo_5809.json",
    "script": "validador_gobernanza_precalculo_5809.py",
}

REQUIRED_SHEETS = [
    "00_DICTAMEN_GLOBAL",
    "01_FUENTES_INVENTARIO",
    "02_HOJAS_INVENTARIO",
    "03_COLUMNAS_INVENTARIO",
    "04_DICCIONARIOS_DETECTADOS",
    "05_DICCIONARIO_ESTADO_PROMEDIOS",
    "06_LLAVES_GOBERNADAS",
    "07_MAPEO_5809_COLS_OFICIALES",
    "08_GOBERNANZA_COLUMNAS_16_19",
    "09_GOBERNANZA_COLUMNAS_20_21",
    "11_TRANSFORMACIONES_AUTORIZ",
    "12_TRANSFORMACIONES_PROHIBIDAS",
    "13_VALIDACIONES_PRE_CALCULO",
    "15_INCIDENTES_METODOLOGICOS",
    "16_BLOQUEOS_ACTUALES",
    "17_REGLAS_DE_REANUDACION",
]

ANNUAL_FIELDS = {
    "16",
    "17",
    "18",
    "19",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
}
ACCUM_FIELDS = {"20", "21", "UNID_CURSADAS_TOTAL", "UNID_APROBADAS_TOTAL"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize(value: Any) -> str:
    text = clean_cell(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    out = []
    for ch in text.upper():
        out.append(ch if ch.isalnum() else "_")
    return "_".join("".join(out).split("_"))


def display(value: Any, max_len: int = 180) -> str:
    text = clean_cell(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def contains_text(df: pd.DataFrame, needle: str) -> bool:
    if df.empty:
        return False
    haystack = " ".join(display(v, 1000) for v in df.astype(str).fillna("").to_numpy().ravel())
    return normalize(needle) in normalize(haystack)


def as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    text = normalize(value)
    return text in {"SI", "TRUE", "VERDADERO", "1", "YES"}


def listify(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [clean_cell(v) for v in value if clean_cell(v)]
    if isinstance(value, tuple):
        return [clean_cell(v) for v in value if clean_cell(v)]
    text = clean_cell(value)
    if not text:
        return []
    if "|" in text:
        return [part.strip() for part in text.split("|") if part.strip()]
    if "," in text:
        return [part.strip() for part in text.split(",") if part.strip()]
    return [text]


def read_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return pd.DataFrame()
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return pd.DataFrame()
        headers = [clean_cell(h) or f"SIN_NOMBRE_{idx + 1}" for idx, h in enumerate(rows[0])]
        records = []
        for row in rows[1:]:
            records.append(
                {
                    headers[idx]: row[idx] if idx < len(row) else ""
                    for idx in range(len(headers))
                }
            )
        return pd.DataFrame(records)
    finally:
        wb.close()


def read_governance(path: Path) -> Tuple[Dict[str, pd.DataFrame], List[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    wb.close()
    data = {sheet: read_sheet(path, sheet) for sheet in sheet_names}
    return data, sheet_names


def safe_col(df: pd.DataFrame, name: str) -> pd.Series:
    if name in df.columns:
        return df[name].fillna("").astype(str)
    return pd.Series([], dtype=str)


def filter_rows_by_column(df: pd.DataFrame, column: str, value: str) -> pd.DataFrame:
    if df.empty or column not in df.columns:
        return pd.DataFrame()
    return df[safe_col(df, column).map(normalize) == normalize(value)]


def any_state_contains(df: pd.DataFrame, state_col: str, expected: str) -> bool:
    if df.empty or state_col not in df.columns:
        return False
    return safe_col(df, state_col).map(lambda x: normalize(expected) in normalize(x)).any()


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    rows = []
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    rows.append("| " + " | ".join(str(col).ljust(widths[idx]) for idx, col in enumerate(cols)) + " |")
    rows.append("| " + " | ".join("-" * width for width in widths) + " |")
    for _, row in df.iterrows():
        rows.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[idx]) for idx, col in enumerate(cols))
            + " |"
        )
    return "\n".join(rows)


def load_plan(plan_path: Path | None) -> Dict[str, Any]:
    if plan_path is None:
        return {
            "descripcion": "Sin plan de calculo/correccion. Validacion estructural de la compuerta previa.",
            "columnas_usadas": [],
            "llaves_usadas": [],
            "diccionarios_usados": [],
            "transformaciones": [],
            "campos_calculo": [],
            "anio_calculo": None,
            "genera_sies_ready": False,
            "modifica_fuentes_originales": False,
            "usa_todos_codcli_rut": False,
            "filtra_por_programa": None,
            "usa_inferencias_como_reglas_oficiales": False,
            "usa_estados_AEIR": False,
            "usa_codcli_lista": False,
            "mezcla_anual_acumulado": False,
            "modo": "SIN_PLAN_OPERATIVO",
        }
    data = json.loads(plan_path.read_text(encoding="utf-8"))
    data.setdefault("descripcion", f"Plan de uso: {plan_path}")
    data.setdefault("columnas_usadas", [])
    data.setdefault("llaves_usadas", [])
    data.setdefault("diccionarios_usados", [])
    data.setdefault("transformaciones", [])
    data.setdefault("campos_calculo", [])
    data.setdefault("anio_calculo", None)
    data.setdefault("genera_sies_ready", False)
    data.setdefault("modifica_fuentes_originales", False)
    data.setdefault("usa_todos_codcli_rut", False)
    data.setdefault("filtra_por_programa", None)
    data.setdefault("usa_inferencias_como_reglas_oficiales", False)
    data.setdefault("usa_estados_AEIR", False)
    data.setdefault("usa_codcli_lista", False)
    data.setdefault("mezcla_anual_acumulado", False)
    data.setdefault("modo", "PLAN_OPERATIVO")
    return data


def build_column_index(columns_df: pd.DataFrame) -> Dict[str, List[Dict[str, str]]]:
    index: Dict[str, List[Dict[str, str]]] = {}
    if columns_df.empty:
        return index
    for _, row in columns_df.iterrows():
        original = clean_cell(row.get("Nombre columna original", ""))
        normalized = clean_cell(row.get("Nombre normalizado", "")) or normalize(original)
        for key in {normalize(original), normalize(normalized)}:
            if not key:
                continue
            index.setdefault(key, []).append({col: clean_cell(row.get(col, "")) for col in columns_df.columns})
    return index


def column_is_governed(column_index: Dict[str, List[Dict[str, str]]], column_name: str) -> Tuple[bool, str]:
    rows = column_index.get(normalize(column_name), [])
    if not rows:
        return False, "Columna no encontrada en inventario de gobernanza."
    states = [row.get("Estado gobernanza", "") for row in rows]
    if any(normalize(state).startswith("GOBERNADA") for state in states):
        return True, "Columna encontrada con estado gobernado en al menos una fuente/hoja."
    return False, "Estados observados: " + "; ".join(sorted(set(states)))


def assess_structural_controls(
    governance: Dict[str, pd.DataFrame],
    sheet_names: List[str],
    fuente_hash_rows: List[Dict[str, Any]],
) -> pd.DataFrame:
    dictamen = governance.get("00_DICTAMEN_GLOBAL", pd.DataFrame())
    columns_df = governance.get("03_COLUMNAS_INVENTARIO", pd.DataFrame())
    dict_estado = governance.get("05_DICCIONARIO_ESTADO_PROMEDIOS", pd.DataFrame())
    keys_df = governance.get("06_LLAVES_GOBERNADAS", pd.DataFrame())
    gov_16_19 = governance.get("08_GOBERNANZA_COLUMNAS_16_19", pd.DataFrame())
    gov_20_21 = governance.get("09_GOBERNANZA_COLUMNAS_20_21", pd.DataFrame())
    authorized = governance.get("11_TRANSFORMACIONES_AUTORIZ", pd.DataFrame())
    prohibited = governance.get("12_TRANSFORMACIONES_PROHIBIDAS", pd.DataFrame())
    blockers = governance.get("16_BLOQUEOS_ACTUALES", pd.DataFrame())
    resume = governance.get("17_REGLAS_DE_REANUDACION", pd.DataFrame())

    n_codcli = filter_rows_by_column(columns_df, "Nombre columna original", "N_CODCLI")
    codcli_lista = filter_rows_by_column(columns_df, "Nombre columna original", "CODCLI_LISTA")
    estado = filter_rows_by_column(columns_df, "Nombre columna original", "ESTADO")
    desc_estado = filter_rows_by_column(columns_df, "Nombre columna original", "DESCRIPCION_ESTADO")
    dict_codes = set(safe_col(dict_estado, "Estado").map(normalize).tolist())
    required_codes = {"A", "E", "I", "R", "NULL_BLANCO"}
    has_null = any("NULL" in code or "BLANCO" in code for code in dict_codes)
    dict_ok = {"A", "E", "I", "R"}.issubset(dict_codes) and has_null

    fuente_hash_ok = all(row["Estado hash"] == "OK" for row in fuente_hash_rows)
    dictamen_no_listo = contains_text(dictamen, DECLARACION_CARGA)

    rows = [
        {
            "ID_CONTROL": "C01",
            "Intento bloqueado": "Usar N_CODCLI como CODCLI academico",
            "Evidencia gobernanza": "N_CODCLI en 03_COLUMNAS_INVENTARIO con BLOQUEADA y prohibicion en 12.",
            "Resultado control": "OK" if any_state_contains(n_codcli, "Estado gobernanza", "BLOQUEADA") and contains_text(prohibited, "N_CODCLI") else "FALTA_CONTROL",
            "Accion": "Bloquear cualquier plan que use N_CODCLI como llave academica.",
        },
        {
            "ID_CONTROL": "C02",
            "Intento bloqueado": "Usar estados A/E/I/R sin diccionario PROMEDIOS",
            "Evidencia gobernanza": "Hoja 05 con A/E/I/R/NULL y columnas ESTADO/DESCRIPCION_ESTADO gobernadas.",
            "Resultado control": "OK" if dict_ok and not estado.empty and not desc_estado.empty else "FALTA_CONTROL",
            "Accion": "Exigir diccionario PROMEDIOS antes de interpretar estados.",
        },
        {
            "ID_CONTROL": "C03",
            "Intento bloqueado": "Calcular 16-19 sin CODCLI_LISTA",
            "Evidencia gobernanza": "CODCLI_LISTA gobernada y hoja 08 documenta columnas 16-19.",
            "Resultado control": "OK" if any_state_contains(codcli_lista, "Estado gobernanza", "GOBERNADA") and not gov_16_19.empty else "FALTA_CONTROL",
            "Accion": "Bloquear calculo 16-19 si CODCLI_LISTA no participa.",
        },
        {
            "ID_CONTROL": "C04",
            "Intento bloqueado": "Sumar todos los CODCLI del RUT sin filtrar por programa",
            "Evidencia gobernanza": "Prohibicion explicita y llave programa en 06_LLAVES_GOBERNADAS.",
            "Resultado control": "OK" if contains_text(prohibited, "Sumar todos los CODCLI") and contains_text(keys_df, "RUT + CODIGO_UNICO + PLAN_ESTUDIOS") else "FALTA_CONTROL",
            "Accion": "Exigir filtro por programa antes de contar unidades.",
        },
        {
            "ID_CONTROL": "C05",
            "Intento bloqueado": "Usar 2026 para calculo anual 2025",
            "Evidencia gobernanza": "Prohibicion explicita y año referencia datos 2025.",
            "Resultado control": "OK" if contains_text(prohibited, "Mezclar 2026") and contains_text(dictamen, "2025") else "FALTA_CONTROL",
            "Accion": "Bloquear ANO=2026 en calculo anual 16-19.",
        },
        {
            "ID_CONTROL": "C06",
            "Intento bloqueado": "Mezclar 16-19 anual con 20-21 acumulado",
            "Evidencia gobernanza": "Hoja 09 separa acumulado; transformacion autorizada obliga separacion.",
            "Resultado control": "OK" if not gov_20_21.empty and contains_text(authorized, "Separar anual 16-19") else "FALTA_CONTROL",
            "Accion": "Bloquear formulas que traten anual y acumulado como el mismo universo.",
        },
        {
            "ID_CONTROL": "C07",
            "Intento bloqueado": "Presentar inferencias como reglas oficiales",
            "Evidencia gobernanza": "Prohibiciones P06/P07 e incidencia metodologica de inferencia A/E/I/R.",
            "Resultado control": "OK" if contains_text(prohibited, "Presentar inferencia") and contains_text(prohibited, "datos observados como norma") else "FALTA_CONTROL",
            "Accion": "Exigir separacion: regla oficial, dato observado, implementacion, decision y pendiente.",
        },
        {
            "ID_CONTROL": "C08",
            "Intento bloqueado": "Generar SIES_READY con bloqueos vigentes",
            "Evidencia gobernanza": "NO_LISTO_PARA_CARGA, bloqueos vigentes y prohibicion SIES_READY.",
            "Resultado control": "OK" if dictamen_no_listo and not blockers.empty and contains_text(prohibited, "SIES_READY") else "FALTA_CONTROL",
            "Accion": "Bloquear cualquier salida final mientras existan bloqueos.",
        },
        {
            "ID_CONTROL": "C09",
            "Intento bloqueado": "Usar columnas no gobernadas",
            "Evidencia gobernanza": "03_COLUMNAS_INVENTARIO contiene Estado gobernanza por columna.",
            "Resultado control": "OK" if not columns_df.empty and "Estado gobernanza" in columns_df.columns else "FALTA_CONTROL",
            "Accion": "Validar cada columna usada contra inventario antes de ejecutar.",
        },
        {
            "ID_CONTROL": "C10",
            "Intento bloqueado": "Modificar fuentes originales",
            "Evidencia gobernanza": "Hashes de fuentes del paso 51 y regla de no modificacion.",
            "Resultado control": "OK" if fuente_hash_ok and contains_text(prohibited, "Modificar fuente original") else "FALTA_CONTROL",
            "Accion": "Comparar hash antes/despues y bloquear escritura sobre fuentes.",
        },
        {
            "ID_CONTROL": "C11",
            "Intento bloqueado": "Reanudar desde paso incorrecto",
            "Evidencia gobernanza": "Reglas de reanudacion: paso 50 reemplaza paso 48.",
            "Resultado control": "OK" if contains_text(resume, "Paso 50 reemplaza paso 48") else "FALTA_CONTROL",
            "Accion": "Bloquear uso de paso 48 como base de decision.",
        },
        {
            "ID_CONTROL": "C12",
            "Intento bloqueado": "Ejecutar con gobernanza incompleta",
            "Evidencia gobernanza": "Hojas minimas requeridas presentes.",
            "Resultado control": "OK" if all(sheet in sheet_names for sheet in REQUIRED_SHEETS) else "FALTA_CONTROL",
            "Accion": "Bloquear si falta alguna hoja critica de gobernanza.",
        },
    ]
    return pd.DataFrame(rows)


def assess_plan(plan: Dict[str, Any], governance: Dict[str, pd.DataFrame]) -> pd.DataFrame:
    columns_df = governance.get("03_COLUMNAS_INVENTARIO", pd.DataFrame())
    blockers = governance.get("16_BLOQUEOS_ACTUALES", pd.DataFrame())
    column_index = build_column_index(columns_df)

    columns_used = listify(plan.get("columnas_usadas"))
    keys_used = listify(plan.get("llaves_usadas"))
    dicts_used = listify(plan.get("diccionarios_usados"))
    transformations = listify(plan.get("transformaciones"))
    calc_fields = {normalize(v) for v in listify(plan.get("campos_calculo"))}
    uses_annual = bool(calc_fields & {normalize(v) for v in ANNUAL_FIELDS})
    uses_accum = bool(calc_fields & {normalize(v) for v in ACCUM_FIELDS})

    rows: List[Dict[str, Any]] = []

    def add(rule_id: str, check: str, triggered: bool, evidence: str, action: str) -> None:
        rows.append(
            {
                "ID_REGLA": rule_id,
                "Validacion": check,
                "Resultado": "BLOQUEADO" if triggered else "OK",
                "Evidencia": evidence,
                "Accion": action if triggered else "Sin bloqueo para el plan evaluado.",
            }
        )

    add(
        "R01",
        "No usar N_CODCLI como CODCLI academico",
        any(normalize(col) == "N_CODCLI" for col in columns_used + keys_used),
        "Columnas/llaves declaradas: " + ", ".join(columns_used + keys_used),
        "Reemplazar por CODCLI_LISTA y filtro por programa.",
    )

    uses_states = as_bool(plan.get("usa_estados_AEIR")) or any(
        normalize(col) in {"ESTADO", "DESCRIPCION_ESTADO"} for col in columns_used
    )
    has_dict = any("DIC_ESTADO_PROMEDIOS" in normalize(item) or "ESTADO_PROMEDIOS" in normalize(item) for item in dicts_used)
    add(
        "R02",
        "No usar A/E/I/R sin diccionario PROMEDIOS",
        uses_states and not has_dict,
        "usa_estados_AEIR="
        + str(plan.get("usa_estados_AEIR"))
        + "; diccionarios_usados="
        + ", ".join(dicts_used),
        "Declarar y aplicar DIC_ESTADO_PROMEDIOS antes de usar ESTADO.",
    )

    has_codcli_lista = as_bool(plan.get("usa_codcli_lista")) or any(
        normalize(col) == "CODCLI_LISTA" for col in columns_used + keys_used
    )
    add(
        "R03",
        "No calcular 16-19 sin CODCLI_LISTA",
        uses_annual and not has_codcli_lista,
        "campos_calculo=" + ", ".join(listify(plan.get("campos_calculo"))) + "; CODCLI_LISTA=" + str(has_codcli_lista),
        "Incluir CODCLI_LISTA desde MAPEO_IDENTIDAD_CODCLI.",
    )

    filtra = plan.get("filtra_por_programa")
    add(
        "R04",
        "No sumar todos los CODCLI del RUT sin filtrar por programa",
        as_bool(plan.get("usa_todos_codcli_rut")) or (uses_annual and filtra is False),
        f"usa_todos_codcli_rut={plan.get('usa_todos_codcli_rut')}; filtra_por_programa={filtra}",
        "Usar llave RUT + CODIGO_UNICO + PLAN_ESTUDIOS y filtro de programa.",
    )

    add(
        "R05",
        "No usar 2026 para calculo anual 2025",
        uses_annual and str(plan.get("anio_calculo")) == "2026",
        f"anio_calculo={plan.get('anio_calculo')}",
        "Cambiar filtro anual a ANO=2025.",
    )

    add(
        "R06",
        "No mezclar 16-19 anual con 20-21 acumulado",
        as_bool(plan.get("mezcla_anual_acumulado")) or (uses_annual and uses_accum),
        "campos_calculo=" + ", ".join(listify(plan.get("campos_calculo"))),
        "Separar formulas, fuentes y validaciones de anual y acumulado.",
    )

    add(
        "R07",
        "No presentar inferencias como reglas oficiales",
        as_bool(plan.get("usa_inferencias_como_reglas_oficiales")),
        f"usa_inferencias_como_reglas_oficiales={plan.get('usa_inferencias_como_reglas_oficiales')}",
        "Mover inferencias a decision interna o pendiente; no regla oficial.",
    )

    add(
        "R08",
        "No generar SIES_READY con bloqueos vigentes",
        as_bool(plan.get("genera_sies_ready")) and not blockers.empty,
        f"genera_sies_ready={plan.get('genera_sies_ready')}; bloqueos_vigentes={len(blockers)}",
        "Mantener salida como no carga hasta cerrar bloqueos.",
    )

    non_governed = []
    for col in columns_used:
        governed, reason = column_is_governed(column_index, col)
        if not governed:
            non_governed.append(f"{col}: {reason}")
    add(
        "R09",
        "No usar columnas no gobernadas",
        bool(non_governed),
        " | ".join(non_governed) if non_governed else "Todas las columnas declaradas estan gobernadas o no se declararon columnas.",
        "Gobernar o retirar columnas antes de ejecutar.",
    )

    add(
        "R10",
        "No modificar fuentes originales",
        as_bool(plan.get("modifica_fuentes_originales")),
        f"modifica_fuentes_originales={plan.get('modifica_fuentes_originales')}",
        "Escribir solo derivados auditados; nunca fuentes originales.",
    )

    if not columns_used and not keys_used and not dicts_used and not transformations and not calc_fields:
        rows.append(
            {
                "ID_REGLA": "R00",
                "Validacion": "Sin plan operativo declarado",
                "Resultado": "OK",
                "Evidencia": "Esta ejecucion valida la compuerta; no habilita calculo ni carga.",
                "Accion": "Para un calculo futuro, entregar plan JSON o ejecutar el validador antes del script operativo.",
            }
        )

    return pd.DataFrame(rows)


def source_hash_control(fuentes_df: pd.DataFrame) -> Tuple[pd.DataFrame, Dict[str, str], Dict[str, str]]:
    rows = []
    before: Dict[str, str] = {}
    registered: Dict[str, str] = {}
    if fuentes_df.empty:
        return pd.DataFrame(), before, registered

    for _, row in fuentes_df.iterrows():
        source_id = clean_cell(row.get("ID_FUENTE", ""))
        path = Path(clean_cell(row.get("Ruta", "")))
        registered_hash = clean_cell(row.get("Hash SHA256", ""))
        registered[source_id] = registered_hash
        if path.exists() and path.is_file():
            current_hash = sha256_file(path)
            before[source_id] = current_hash
            state = "OK" if current_hash == registered_hash else "DIFIERE_DE_GOBERNANZA"
        else:
            current_hash = ""
            state = "NO_EXISTE"
        rows.append(
            {
                "ID_FUENTE": source_id,
                "Ruta": str(path),
                "Hash gobernanza 51": registered_hash,
                "Hash actual previo": current_hash,
                "Estado hash": state,
                "Puede modificar regla oficial": clean_cell(row.get("Puede modificar regla oficial", "")),
                "Observacion": clean_cell(row.get("Observacion", "")),
            }
        )
    return pd.DataFrame(rows), before, registered


def source_hash_after(fuentes_df: pd.DataFrame) -> Dict[str, str]:
    hashes: Dict[str, str] = {}
    if fuentes_df.empty:
        return hashes
    for _, row in fuentes_df.iterrows():
        source_id = clean_cell(row.get("ID_FUENTE", ""))
        path = Path(clean_cell(row.get("Ruta", "")))
        if path.exists() and path.is_file():
            hashes[source_id] = sha256_file(path)
        else:
            hashes[source_id] = ""
    return hashes


def build_critical_columns(governance: Dict[str, pd.DataFrame]) -> pd.DataFrame:
    columns_df = governance.get("03_COLUMNAS_INVENTARIO", pd.DataFrame())
    if columns_df.empty:
        return pd.DataFrame()
    critical = {
        "N_CODCLI",
        "CODCLI_LISTA",
        "ESTADO",
        "DESCRIPCION_ESTADO",
        "CODCLI",
        "ANO",
        "PERIODO",
        "CODRAMO",
        "CURSO_1ER_SEM",
        "CURSO_2DO_SEM",
        "UNIDADES_CURSADAS",
        "UNIDADES_APROBADAS",
        "UNID_CURSADAS_TOTAL",
        "UNID_APROBADAS_TOTAL",
    }
    mask = safe_col(columns_df, "Nombre columna original").map(normalize).isin({normalize(c) for c in critical})
    cols = [
        "Archivo",
        "Hoja",
        "Nombre columna original",
        "Nombre normalizado",
        "Posible rol",
        "Uso permitido",
        "Uso prohibido",
        "Estado gobernanza",
        "Justificacion",
    ]
    return columns_df.loc[mask, [c for c in cols if c in columns_df.columns]].copy()


def build_keys_dict_trans(governance: Dict[str, pd.DataFrame]) -> pd.DataFrame:
    frames = []
    for sheet, label in [
        ("06_LLAVES_GOBERNADAS", "LLAVE"),
        ("04_DICCIONARIOS_DETECTADOS", "DICCIONARIO"),
        ("11_TRANSFORMACIONES_AUTORIZ", "TRANSFORMACION_AUTORIZADA"),
        ("12_TRANSFORMACIONES_PROHIBIDAS", "TRANSFORMACION_PROHIBIDA"),
    ]:
        df = governance.get(sheet, pd.DataFrame()).copy()
        if df.empty:
            continue
        df.insert(0, "Tipo control", label)
        df.insert(1, "Hoja gobernanza", sheet)
        frames.append(df)
    if not frames:
        return pd.DataFrame()
    out = pd.concat(frames, ignore_index=True, sort=False)
    preferred = [
        "Tipo control",
        "Hoja gobernanza",
        "ID_LLAVE",
        "Nombre llave",
        "ID_DICCIONARIO",
        "Campo codigo",
        "Campo descripcion",
        "Codigo",
        "Descripcion",
        "ID_TRANSFORMACION",
        "Descripcion",
        "ID_PROHIBICION",
        "Transformacion prohibida",
        "Estado",
        "Estado gobernanza",
        "Evidencia",
        "Observacion",
    ]
    cols = [c for c in preferred if c in out.columns] + [c for c in out.columns if c not in preferred]
    return out[cols]


def build_manifest_readable(manifest: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    for key in [
        "proceso",
        "subproyecto",
        "timestamp",
        "estado_carga",
        "estado_validador",
        "fuentes_originales_modificadas",
        "gobernanza_xlsx",
        "plan_uso",
    ]:
        rows.append({"Seccion": "GLOBAL", "Clave": key, "Valor": manifest.get(key, "")})
    for key, value in manifest.get("conteos", {}).items():
        rows.append({"Seccion": "CONTEOS", "Clave": key, "Valor": value})
    for key, value in manifest.get("archivos", {}).items():
        rows.append({"Seccion": "ARCHIVOS", "Clave": key, "Valor": value})
    return pd.DataFrame(rows)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="7030A0")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, min(ws.max_column, 45) + 1):
            letter = get_column_letter(col_idx)
            max_len = 12
            for cell in ws[letter][: min(ws.max_row, 200)]:
                text = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(text), 90))
            ws.column_dimensions[letter].width = min(max_len + 2, 65)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    wb.close()


def write_excel(path: Path, dataframes: OrderedDict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, df in dataframes.items():
            safe_df = df.copy()
            for col in safe_df.columns:
                safe_df[col] = safe_df[col].map(lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
            safe_df.to_excel(writer, sheet_name=sheet_name, index=False)
    style_workbook(path)


def build_markdown(
    dictamen_df: pd.DataFrame,
    controls_df: pd.DataFrame,
    plan_df: pd.DataFrame,
    blockers_df: pd.DataFrame,
    manifest: Dict[str, Any],
) -> str:
    lines = [
        "# Validacion gobernanza precalculo 5809",
        "",
        f"Proceso: **{PROCESO}**",
        f"Subproyecto: **{SUBPROYECTO}**",
        f"Año proceso: **{ANIO_PROCESO}**",
        f"Año referencia datos: **{ANIO_REFERENCIA_DATOS}**",
        f"Declaración carga: **{DECLARACION_CARGA}**",
        f"Estado validador: **{manifest['estado_validador']}**",
        "",
        "## Dictamen",
        "",
        markdown_table(dictamen_df),
        "",
        "## Controles obligatorios",
        "",
        markdown_table(controls_df),
        "",
        "## Validacion del plan de uso",
        "",
        markdown_table(plan_df),
        "",
        "## Bloqueos vigentes heredados de gobernanza 51",
        "",
        markdown_table(blockers_df),
        "",
        "## Decision",
        "",
        "- Este paso no corrige datos, no recalcula columnas finales y no genera SIES_READY.",
        "- Si un plan futuro activa cualquiera de las reglas bloqueantes, el validador debe devolver BLOQUEADO.",
        "- Aunque el validador quede apto como compuerta, la carga sigue NO_LISTO_PARA_CARGA mientras existan bloqueos.",
        "",
        "## Archivos",
        "",
    ]
    for key, value in manifest.get("archivos", {}).items():
        lines.append(f"- {key}: `{value}`")
    lines.append("")
    return "\n".join(lines)


def scan_sies_ready(output_dir: Path) -> List[str]:
    return [str(path) for path in output_dir.rglob("*SIES_READY*")]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--timestamp", default=datetime.now().strftime("%Y%m%d_%H%M%S"))
    parser.add_argument("--gobernanza", default=str(GOBERNANZA_XLSX))
    parser.add_argument("--plan-json", default=None, help="Plan futuro a validar. Si se omite, valida solo la compuerta.")
    args = parser.parse_args()

    output_dir = BASE_52 / f"VALIDADOR_GOBERNANZA_PRECALCULO_5809_{args.timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = output_dir / OUTPUT_FILES["excel"]
    md_path = output_dir / OUTPUT_FILES["informe"]
    manifest_path = output_dir / OUTPUT_FILES["manifest"]
    script_path = output_dir / OUTPUT_FILES["script"]
    current_script = Path(__file__).resolve()
    if current_script != script_path.resolve():
        shutil.copy2(current_script, script_path)

    governance_path = Path(args.gobernanza)
    if not governance_path.exists():
        raise FileNotFoundError(governance_path)

    plan_path = Path(args.plan_json) if args.plan_json else None
    plan = load_plan(plan_path)

    governance, sheet_names = read_governance(governance_path)
    fuentes_df = governance.get("01_FUENTES_INVENTARIO", pd.DataFrame())
    hash_df, source_hashes_before, source_hashes_registered = source_hash_control(fuentes_df)
    fuente_hash_rows = hash_df.to_dict(orient="records")

    controls_df = assess_structural_controls(governance, sheet_names, fuente_hash_rows)
    plan_df = assess_plan(plan, governance)
    blockers_df = governance.get("16_BLOQUEOS_ACTUALES", pd.DataFrame()).copy()
    critical_columns_df = build_critical_columns(governance)
    keys_dict_trans_df = build_keys_dict_trans(governance)
    governance_read_df = pd.DataFrame(
        [
            {
                "Hoja": sheet,
                "Filas": len(governance.get(sheet, pd.DataFrame())),
                "Columnas": len(governance.get(sheet, pd.DataFrame()).columns),
                "Requerida": "SI" if sheet in REQUIRED_SHEETS else "NO",
                "Estado": "OK" if sheet in sheet_names else "FALTA",
            }
            for sheet in sheet_names
        ]
    )

    structural_blocked = bool((controls_df["Resultado control"] != "OK").any())
    plan_blocked = bool((plan_df["Resultado"] == "BLOQUEADO").any())
    hash_blocked = bool((hash_df["Estado hash"] != "OK").any()) if not hash_df.empty else True
    estado_validador = ESTADO_BLOQUEADO if structural_blocked or plan_blocked or hash_blocked else ESTADO_APTO

    sources_after_pre_outputs = source_hash_after(fuentes_df)
    fuentes_modificadas = "SI" if source_hashes_before != sources_after_pre_outputs else "NO"
    if fuentes_modificadas == "SI":
        estado_validador = ESTADO_BLOQUEADO

    sies_ready_found_pre = scan_sies_ready(output_dir)
    if sies_ready_found_pre:
        estado_validador = ESTADO_BLOQUEADO

    dictamen_df = pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Gobernanza leida": str(governance_path),
                "Declaracion carga": DECLARACION_CARGA,
                "Estado validador": estado_validador,
                "Fuentes originales modificadas": fuentes_modificadas,
                "Hojas gobernanza leidas": len(sheet_names),
                "Controles obligatorios": len(controls_df),
                "Controles faltantes": int((controls_df["Resultado control"] != "OK").sum()),
                "Reglas plan bloqueadas": int((plan_df["Resultado"] == "BLOQUEADO").sum()),
                "Bloqueos vigentes heredados": len(blockers_df),
                "SIES_READY generado": "SI" if sies_ready_found_pre else "NO",
                "Dictamen": (
                    "Validador apto como control previo. No habilita carga ni calculo definitivo."
                    if estado_validador == ESTADO_APTO
                    else "Validador bloqueado: revisar controles, plan de uso o hashes de fuente."
                ),
                "Proximo paso permitido": (
                    "Solo validar un plan tecnico futuro; no corregir datos ni generar carga."
                ),
            }
        ]
    )

    manifest: Dict[str, Any] = {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": args.timestamp,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "estado_carga": DECLARACION_CARGA,
        "estado_validador": estado_validador,
        "fuentes_originales_modificadas": fuentes_modificadas,
        "gobernanza_xlsx": str(governance_path),
        "gobernanza_hash_sha256": sha256_file(governance_path),
        "plan_uso": str(plan_path) if plan_path else "SIN_PLAN_OPERATIVO",
        "plan": plan,
        "conteos": {
            "hojas_gobernanza_leidas": len(sheet_names),
            "controles_obligatorios": len(controls_df),
            "controles_faltantes": int((controls_df["Resultado control"] != "OK").sum()),
            "reglas_plan_bloqueadas": int((plan_df["Resultado"] == "BLOQUEADO").sum()),
            "bloqueos_vigentes_heredados": len(blockers_df),
            "fuentes_hash_ok": int((hash_df["Estado hash"] == "OK").sum()) if not hash_df.empty else 0,
            "fuentes_hash_difieren": int((hash_df["Estado hash"] != "OK").sum()) if not hash_df.empty else 0,
        },
        "controles": controls_df.to_dict(orient="records"),
        "validacion_plan": plan_df.to_dict(orient="records"),
        "bloqueos_vigentes": blockers_df.to_dict(orient="records"),
        "hashes_fuentes_registrados": source_hashes_registered,
        "hashes_fuentes_antes": source_hashes_before,
        "hashes_fuentes_despues_pre_outputs": sources_after_pre_outputs,
        "sies_ready_detectado": sies_ready_found_pre,
        "archivos": {
            "Excel": str(excel_path),
            "Informe": str(md_path),
            "Manifest": str(manifest_path),
            "Script": str(script_path),
        },
    }
    manifest_readable_df = build_manifest_readable(manifest)
    rules_block_df = pd.DataFrame(
        [
            {
                "ID": "B01",
                "Regla bloqueo": "Se intenta usar N_CODCLI como CODCLI academico.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B02",
                "Regla bloqueo": "Se intenta usar estados A/E/I/R sin diccionario PROMEDIOS.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B03",
                "Regla bloqueo": "Se intenta calcular 16-19 sin CODCLI_LISTA.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B04",
                "Regla bloqueo": "Se intenta sumar todos los CODCLI del RUT sin filtrar por programa.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B05",
                "Regla bloqueo": "Se intenta usar 2026 para calculo anual 2025.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B06",
                "Regla bloqueo": "Se intenta mezclar 16-19 anual con 20-21 acumulado.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B07",
                "Regla bloqueo": "Se intenta presentar inferencias como reglas oficiales.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B08",
                "Regla bloqueo": "Se intenta generar SIES_READY con bloqueos vigentes.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B09",
                "Regla bloqueo": "Se intenta usar columnas no gobernadas.",
                "Resultado esperado": "BLOQUEADO",
            },
            {
                "ID": "B10",
                "Regla bloqueo": "Se intenta modificar fuentes originales.",
                "Resultado esperado": "BLOQUEADO",
            },
        ]
    )

    dataframes = OrderedDict(
        [
            ("00_DICTAMEN_VALIDADOR", dictamen_df),
            ("01_GOBERNANZA_LEIDA", governance_read_df),
            ("02_CONTROLES_OBLIGATORIOS", controls_df),
            ("03_VALIDACION_PLAN_USO", plan_df),
            ("04_COLUMNAS_CRITICAS", critical_columns_df),
            ("05_LLAVES_DICC_TRANSF", keys_dict_trans_df),
            ("06_BLOQUEOS_VIGENTES", blockers_df),
            ("07_HASHES_FUENTES", hash_df),
            ("08_REGLAS_BLOQUEO", rules_block_df),
            ("09_MANIFEST_LEGIBLE", manifest_readable_df),
        ]
    )
    write_excel(excel_path, dataframes)

    md_text = build_markdown(dictamen_df, controls_df, plan_df, blockers_df, manifest)
    md_path.write_text(md_text, encoding="utf-8")

    sources_after = source_hash_after(fuentes_df)
    manifest["hashes_fuentes_despues"] = sources_after
    if source_hashes_before != sources_after:
        manifest["fuentes_originales_modificadas"] = "SI"
        manifest["estado_validador"] = ESTADO_BLOQUEADO
        dictamen_df.loc[0, "Fuentes originales modificadas"] = "SI"
        dictamen_df.loc[0, "Estado validador"] = ESTADO_BLOQUEADO

    sies_ready_found = scan_sies_ready(output_dir)
    manifest["sies_ready_detectado"] = sies_ready_found
    if sies_ready_found:
        manifest["estado_validador"] = ESTADO_BLOQUEADO
        dictamen_df.loc[0, "SIES_READY generado"] = "SI"
        dictamen_df.loc[0, "Estado validador"] = ESTADO_BLOQUEADO

    manifest["hashes_salidas"] = {
        "Excel": sha256_file(excel_path),
        "Informe": sha256_file(md_path),
        "Script": sha256_file(script_path),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print("VALIDADOR GOBERNANZA PRECALCULO 5809 — GENERADO")
    print(f"Fuentes originales modificadas: {manifest['fuentes_originales_modificadas']}")
    print(f"Declaración carga: {DECLARACION_CARGA}")
    print(f"Estado validador: {manifest['estado_validador']}")
    print()
    print("DICTAMEN")
    print(markdown_table(dictamen_df))
    print()
    print("CONTROLES OBLIGATORIOS")
    print(markdown_table(controls_df[["ID_CONTROL", "Intento bloqueado", "Resultado control"]]))
    print()
    print("ARCHIVOS")
    print(f"Excel: {excel_path}")
    print(f"Informe: {md_path}")
    print(f"Manifest: {manifest_path}")
    print(f"Script: {script_path}")


if __name__ == "__main__":
    main()
