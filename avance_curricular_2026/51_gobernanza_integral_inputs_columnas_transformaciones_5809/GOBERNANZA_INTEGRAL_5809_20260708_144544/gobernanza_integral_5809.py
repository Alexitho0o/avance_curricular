#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gobernanza integral de inputs, columnas, llaves, diccionarios,
transformaciones y decisiones para el archivo 5809.

Este script no modifica fuentes originales, no genera archivos de carga
SIES_READY y no recalcula columnas definitivas. Su salida es un expediente de
gobernanza auditable para decidir la reanudación del proceso.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = (
    "Gobernanza integral de inputs, raw, columnas, llaves, diccionarios, "
    "transformaciones y trazabilidad para archivo 5809 Matrícula Avance Curricular"
)
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_51 = REPO / "avance_curricular_2026" / "51_gobernanza_integral_inputs_columnas_transformaciones_5809"

PATHS = OrderedDict(
    [
        (
            "F01_INSTRUCTIVO_OFICIAL",
            {
                "ruta": REPO
                / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/Instructivo_Avance Curricular SIES - 2026.txt",
                "clasificacion": "instructivo",
                "origen": "SIES / Subsecretaría de Educación Superior",
                "nivel": "A",
                "regla": "SI",
                "dato": "SI",
                "implementacion": "NO",
                "estado": "vigente",
                "observacion": "Fuente normativa oficial del proceso. No se modifica.",
            },
        ),
        (
            "F02_PRECARGA_5809",
            {
                "ruta": REPO
                / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/5809_Precarga Matrícula Avance Curricular 2026.csv",
                "clasificacion": "precarga",
                "origen": "SIES / precarga matrícula avance curricular",
                "nivel": "A",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "NO",
                "estado": "vigente",
                "observacion": "Base oficial 5809. Debe permanecer intacta.",
            },
        ),
        (
            "F03_PRECARGA_5810",
            {
                "ruta": REPO
                / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/5810_Precarga Carreras Avance Curricular 20268.csv",
                "clasificacion": "precarga",
                "origen": "SIES / precarga carreras avance curricular",
                "nivel": "A",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "NO",
                "estado": "vigente",
                "observacion": "Referencia oficial de carreras/planes cargables. Debe permanecer intacta.",
            },
        ),
        (
            "F04_PROMEDIOS_ACTUALIZADO",
            {
                "ruta": REPO
                / "avance_curricular_2026/25_actualizacion_fuente_promedios/CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/00_FUENTE_CONGELADA/PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx",
                "clasificacion": "fuente institucional",
                "origen": "Fuente académica institucional PROMEDIOSDEALUMNOS",
                "nivel": "B",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "SI",
                "estado": "vigente",
                "observacion": "Fuente académica primaria para detalle de ramos, año, periodo y estado.",
            },
        ),
        (
            "F05_MAPEO_IDENTIDAD_CODCLI",
            {
                "ruta": REPO
                / "avance_curricular_2026/05_cierre_integral/CIERRE_AVANCE_CURRICULAR_20260703_125157/04_CONCILIACION_MATRICULA/PREPARACION_MATRICULA_20260703_141254/02_IDENTIFICADORES/MAPEO_IDENTIDAD_CODCLI.xlsx",
                "clasificacion": "mapeo",
                "origen": "Conciliación institucional 5809-RUT-CODCLI",
                "nivel": "B",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "SI",
                "estado": "vigente",
                "observacion": "Fuente de identidad; CODCLI_LISTA es la columna gobernada para CODCLI académico.",
            },
        ),
        (
            "F06_PLANES_ESTUDIO",
            {
                "ruta": REPO
                / "avance_curricular_2026/01_fuentes_institucionales/planes_estudio/Listado_Planes_estudio_Sedes_RE_CO_20260701.xlsx",
                "clasificacion": "fuente institucional",
                "origen": "Planes de estudio institucionales",
                "nivel": "B",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "SI",
                "estado": "vigente",
                "observacion": "Referencia de planes, ramos y sedes; no reemplaza reglas oficiales.",
            },
        ),
        (
            "F07_RECALCULO_50_EVIDENCIA",
            {
                "ruta": REPO
                / "avance_curricular_2026/50_recalculo_16_19_diccionario_estado_promedios/RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS_20260708_143748/02_RESULTADOS/RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS.xlsx",
                "clasificacion": "evidencia",
                "origen": "Paso 50 / recálculo con diccionario estado PROMEDIOS",
                "nivel": "C",
                "regla": "NO",
                "dato": "SI",
                "implementacion": "SI",
                "estado": "derivado",
                "observacion": "Evidencia técnica no normativa. Reemplaza inferencias previas sobre estados.",
            },
        ),
    ]
)

REQUESTED_SHEETS = OrderedDict(
    [
        ("00_DICTAMEN_GLOBAL", "00_DICTAMEN_GLOBAL"),
        ("01_FUENTES_INVENTARIO", "01_FUENTES_INVENTARIO"),
        ("02_HOJAS_INVENTARIO", "02_HOJAS_INVENTARIO"),
        ("03_COLUMNAS_INVENTARIO", "03_COLUMNAS_INVENTARIO"),
        ("04_DICCIONARIOS_DETECTADOS", "04_DICCIONARIOS_DETECTADOS"),
        ("05_DICCIONARIO_ESTADO_PROMEDIOS", "05_DICCIONARIO_ESTADO_PROMEDIOS"),
        ("06_LLAVES_GOBERNADAS", "06_LLAVES_GOBERNADAS"),
        ("07_MAPEO_5809_COLUMNAS_OFICIALES", "07_MAPEO_5809_COLS_OFICIALES"),
        ("08_GOBERNANZA_COLUMNAS_16_19", "08_GOBERNANZA_COLUMNAS_16_19"),
        ("09_GOBERNANZA_COLUMNAS_20_21", "09_GOBERNANZA_COLUMNAS_20_21"),
        ("10_FLUJO_RAW_A_PROCESADO", "10_FLUJO_RAW_A_PROCESADO"),
        ("11_TRANSFORMACIONES_AUTORIZADAS", "11_TRANSFORMACIONES_AUTORIZ"),
        ("12_TRANSFORMACIONES_PROHIBIDAS", "12_TRANSFORMACIONES_PROHIBIDAS"),
        ("13_VALIDACIONES_OBLIGATORIAS_PRE_CALCULO", "13_VALIDACIONES_PRE_CALCULO"),
        ("14_VALIDACIONES_OBLIGATORIAS_POST_CALCULO", "14_VALIDACIONES_POST_CALCULO"),
        ("15_INCIDENTES_METODOLOGICOS", "15_INCIDENTES_METODOLOGICOS"),
        ("16_BLOQUEOS_ACTUALES", "16_BLOQUEOS_ACTUALES"),
        ("17_REGLAS_DE_REANUDACION", "17_REGLAS_DE_REANUDACION"),
        ("18_RESUMEN_EJECUTIVO", "18_RESUMEN_EJECUTIVO"),
        ("19_FUENTES", "19_FUENTES"),
        ("20_MANIFEST_LEGIBLE", "20_MANIFEST_LEGIBLE"),
    ]
)

OFFICIAL_5809_COLUMNS = [
    "CODIGO_IES_NUM",
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "CURSO_1ER_SEM",
    "CURSO_2DO_SEM",
    "UNIDADES_CURSADAS",
    "UNIDADES_APROBADAS",
    "UNID_CURSADAS_TOTAL",
    "UNID_APROBADAS_TOTAL",
    "VIGENCIA",
]

DICT_ESTADO_PROMEDIOS = [
    {
        "Estado": "A",
        "Descripcion": "APROBADO",
        "Dato observado en PROMEDIOS": "ESTADO=A; DESCRIPCION_ESTADO=APROBADO",
        "Regla funcional aplicable": "Actividad con resultado aprobatorio formal.",
        "Tratamiento tecnico": "Suma como aprobada.",
        "Pendiente de confirmacion": "NO para lectura del diccionario; validar que el ramo pertenece al programa 2025.",
        "Nivel respaldo": "B",
    },
    {
        "Estado": "E",
        "Descripcion": "CONVALIDACION",
        "Dato observado en PROMEDIOS": "ESTADO=E; DESCRIPCION_ESTADO=CONVALIDACION",
        "Regla funcional aplicable": (
            "Dato observado como aprobatorio/convalidado; el instructivo excluye validaciones "
            "en aprobadas anuales 2025 y las permite en acumuladas."
        ),
        "Tratamiento tecnico": (
            "Suma como aprobada solo si el campo a calcular permite convalidaciones o si "
            "validacion funcional confirma que la fuente la reporta como actividad aprobada aplicable."
        ),
        "Pendiente de confirmacion": "SI para uso definitivo en UNIDADES_APROBADAS anual.",
        "Nivel respaldo": "B",
    },
    {
        "Estado": "I",
        "Descripcion": "HOMOLOGADO",
        "Dato observado en PROMEDIOS": "ESTADO=I; DESCRIPCION_ESTADO=HOMOLOGADO",
        "Regla funcional aplicable": (
            "Dato observado como aprobatorio/homologado; requiere separar anual 2025 de acumulado."
        ),
        "Tratamiento tecnico": "Suma como aprobada/homologada, sujeto a validacion funcional por campo.",
        "Pendiente de confirmacion": "SI para uso definitivo en UNIDADES_APROBADAS anual.",
        "Nivel respaldo": "B",
    },
    {
        "Estado": "R",
        "Descripcion": "REPROBADO",
        "Dato observado en PROMEDIOS": "ESTADO=R; DESCRIPCION_ESTADO=REPROBADO",
        "Regla funcional aplicable": "Actividad con resultado no aprobatorio.",
        "Tratamiento tecnico": "No suma como aprobada; puede contar como cursada si pertenece a 2025/programa.",
        "Pendiente de confirmacion": "NO para lectura del diccionario; validar pertenencia del ramo al programa.",
        "Nivel respaldo": "B",
    },
    {
        "Estado": "NULL / blanco",
        "Descripcion": "(en blanco)",
        "Dato observado en PROMEDIOS": "Estado nulo/blanco observado en evidencia del paso 50.",
        "Regla funcional aplicable": "No existe regla explicita cerrada para computarlo automaticamente.",
        "Tratamiento tecnico": "No se aplica automaticamente; queda en revision.",
        "Pendiente de confirmacion": "SI.",
        "Nivel respaldo": "B",
    },
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    cleaned = []
    for ch in text.upper():
        if ch.isalnum():
            cleaned.append(ch)
        else:
            cleaned.append("_")
    out = "_".join("".join(cleaned).split("_"))
    return out


def display_value(value: Any, max_len: int = 120) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def is_empty(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def infer_atomic_type(value: Any) -> str:
    if is_empty(value):
        return "vacio"
    if isinstance(value, bool):
        return "booleano"
    if isinstance(value, int):
        return "entero"
    if isinstance(value, float):
        return "decimal"
    if hasattr(value, "isoformat") and value.__class__.__name__ in {"date", "datetime", "Timestamp"}:
        return "fecha"
    text = str(value).strip()
    if text.upper() in {"SI", "NO", "SÍ"}:
        return "booleano_texto"
    try:
        int(text)
        return "entero_texto"
    except Exception:
        pass
    try:
        float(text.replace(",", "."))
        return "decimal_texto"
    except Exception:
        pass
    return "texto"


def summarize_type(types: Iterable[str]) -> str:
    clean = sorted(t for t in set(types) if t != "vacio")
    if not clean:
        return "vacio"
    numeric = {"entero", "decimal", "entero_texto", "decimal_texto"}
    if set(clean).issubset(numeric):
        return "numerico"
    if len(clean) == 1:
        return clean[0]
    return "mixto: " + ", ".join(clean)


def detect_csv(path: Path) -> Tuple[str, str]:
    raw = path.read_bytes()[:8192]
    encoding = "latin1"
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
        try:
            sample = raw.decode(enc)
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    sample = raw.decode(encoding, errors="replace")
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t,")
        sep = dialect.delimiter
    except Exception:
        sep = ";"
    return encoding, sep


def update_stats(stats: List[Dict[str, Any]], row: List[Any]) -> None:
    if len(row) < len(stats):
        row = row + [None] * (len(stats) - len(row))
    for idx, value in enumerate(row[: len(stats)]):
        item = stats[idx]
        item["tipos"].add(infer_atomic_type(value))
        if is_empty(value):
            continue
        item["no_vacios"] += 1
        text = display_value(value, 120)
        if len(item["distintos_set"]) < 50000:
            item["distintos_set"].add(text)
        else:
            item["distintos_capped"] = True
        if text not in item["muestras"] and len(item["muestras"]) < 8:
            item["muestras"].append(text)


def make_stats(headers: List[Any]) -> List[Dict[str, Any]]:
    normalized_headers = []
    for pos, header in enumerate(headers, start=1):
        original = display_value(header)
        if not original:
            original = f"SIN_NOMBRE_COL_{pos}"
        normalized_headers.append(original)
    return [
        {
            "original": header,
            "normalizado": normalize_text(header),
            "posicion": pos,
            "no_vacios": 0,
            "distintos_set": set(),
            "distintos_capped": False,
            "muestras": [],
            "tipos": set(),
        }
        for pos, header in enumerate(normalized_headers, start=1)
    ]


def finish_stats(stats: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    out = []
    for item in stats:
        distintos = (
            f">{len(item['distintos_set'])}"
            if item["distintos_capped"]
            else str(len(item["distintos_set"]))
        )
        out.append(
            {
                "original": item["original"],
                "normalizado": item["normalizado"],
                "posicion": item["posicion"],
                "tipo_observado": summarize_type(item["tipos"]),
                "no_vacios": item["no_vacios"],
                "distintos": distintos,
                "muestra": " | ".join(item["muestras"]),
            }
        )
    return out


def classify_sheet(source_id: str, sheet_name: str) -> Dict[str, str]:
    normalized = normalize_text(sheet_name)
    is_recalc = source_id == "F07_RECALCULO_50_EVIDENCIA"
    is_prom = source_id == "F04_PROMEDIOS_ACTUALIZADO"
    is_map = source_id == "F05_MAPEO_IDENTIDAD_CODCLI"
    is_precarga = source_id in {"F02_PRECARGA_5809", "F03_PRECARGA_5810"}
    is_dict = normalized in {"01_DICCIONARIO_ESTADO"} or (
        is_prom and sheet_name == "Hoja1"
    )
    if source_id == "F01_INSTRUCTIVO_OFICIAL":
        uso_candidato = "Regla oficial y validaciones oficiales."
        uso_permitido = "Definir reglas oficiales, campos obligatorios y restricciones."
        uso_prohibido = "No usar como dato observado individual."
        estado = "GOBERNADA"
    elif is_recalc:
        uso_candidato = "Evidencia tecnica posterior al error metodologico."
        uso_permitido = "Auditar bloqueos, impacto y reemplazo de inferencias previas."
        uso_prohibido = "No usar como regla oficial ni como base de carga definitiva."
        estado = "GOBERNADA_COMO_EVIDENCIA"
    elif is_map:
        uso_candidato = "Mapeo de identidad y CODCLI_LISTA."
        uso_permitido = "Cruce controlado entre 5809, RUT y CODCLI academico."
        uso_prohibido = "No usar N_CODCLI como CODCLI academico."
        estado = "GOBERNADA_CON_RESTRICCIONES"
    elif is_prom:
        uso_candidato = "Detalle academico y diccionario de estados."
        uso_permitido = "Dato observado para ramos, ANO, PERIODO, ESTADO, CODCLI."
        uso_prohibido = "No imponer regla oficial ni mezclar 2026 con calculo anual 2025."
        estado = "GOBERNADA_PARCIAL"
    elif is_precarga:
        uso_candidato = "Precarga oficial raw."
        uso_permitido = "Base de identificacion y orden oficial."
        uso_prohibido = "No modificar fuente original."
        estado = "GOBERNADA_RAW"
    else:
        uso_candidato = "Referencia institucional."
        uso_permitido = "Referencia y validacion auxiliar."
        uso_prohibido = "No reemplazar reglas oficiales."
        estado = "PENDIENTE_VALIDACION_FUNCIONAL"
    return {
        "Uso candidato": uso_candidato,
        "Uso permitido": uso_permitido,
        "Uso prohibido": uso_prohibido,
        "Es fuente raw": "SI" if source_id in {"F01_INSTRUCTIVO_OFICIAL", "F02_PRECARGA_5809", "F03_PRECARGA_5810", "F04_PROMEDIOS_ACTUALIZADO", "F05_MAPEO_IDENTIDAD_CODCLI", "F06_PLANES_ESTUDIO"} else "NO",
        "Es derivado": "SI" if is_recalc else "NO",
        "Es auditoria": "SI" if is_recalc or "RESUMEN" in normalized else "NO",
        "Es diccionario": "SI" if is_dict else "NO",
        "Es detalle academico": "SI" if is_prom and sheet_name == "Hoja1" else "NO",
        "Es precarga": "SI" if is_precarga else "NO",
        "Es mapeo": "SI" if is_map else "NO",
        "Estado gobernanza": estado,
        "Observacion": "Hoja clasificada por fuente, uso permitido y restricciones de gobernanza.",
    }


def classify_column(
    source_id: str,
    sheet_name: str,
    original: str,
    normalized: str,
    position: int,
) -> Tuple[str, str, str, str, str, str]:
    role = "atributo"
    uso_permitido = "Solo inventario hasta que exista decision explicita."
    uso_prohibido = "No usar para calculo definitivo sin gobernanza."
    nivel = PATHS[source_id]["nivel"]
    estado = "NO_GOBERNADA"
    justificacion = "Columna inventariada; no se requiere para el calculo controlado actual o no tiene regla cerrada."

    if normalized.startswith("SIN_NOMBRE"):
        return (
            "auditoria",
            "Solo evidenciar estructura vacia o columna sin encabezado.",
            "No usar en cruces ni calculos.",
            nivel,
            "BLOQUEADA",
            "Encabezado vacio o no gobernable.",
        )

    if normalized in {"CODIGO_UNICO", "PLAN_ESTUDIOS", "NUM_DOCUMENTO", "DV"}:
        role = "llave"
        uso_permitido = "Parte de llave oficial 5809 o llave de programa."
        uso_prohibido = "No alterar si proviene de precarga oficial; no resolver identidad por nombre."
        estado = "GOBERNADA"
        justificacion = "Llave oficial/documental requerida para trazabilidad 5809."

    if normalized in {"CODCLI", "CODCLI_NORM"}:
        role = "llave"
        uso_permitido = "Llave academica en PROMEDIOS o evidencia tecnica, una vez conciliada."
        uso_prohibido = "No usar sin vincular a RUT/programa cuando hay multiples CODCLI."
        estado = "GOBERNADA" if source_id in {"F04_PROMEDIOS_ACTUALIZADO", "F07_RECALCULO_50_EVIDENCIA"} else "PENDIENTE"
        justificacion = "CODCLI es identificador academico observado; requiere filtro por programa."

    if normalized == "CODCLI_LISTA":
        role = "llave"
        uso_permitido = "Columna correcta de mapeo para listar CODCLI academicos conciliados."
        uso_prohibido = "No sumar todos los CODCLI sin filtrar por programa y año."
        estado = "GOBERNADA"
        justificacion = "La gobernanza corrige el error de usar N_CODCLI como CODCLI academico."

    if normalized == "N_CODCLI":
        return (
            "auditoria",
            "Conteo de CODCLI asociados a identidad; solo sirve para clasificar multiplicidad.",
            "PROHIBIDO usar como CODCLI academico o llave de detalle.",
            nivel,
            "BLOQUEADA",
            "Incidente metodologico: N_CODCLI es conteo, no identificador academico.",
        )

    if normalized in {"ESTADO", "DESCRIPCION_ESTADO"} and source_id in {
        "F04_PROMEDIOS_ACTUALIZADO",
        "F07_RECALCULO_50_EVIDENCIA",
    }:
        role = "diccionario"
        uso_permitido = "Diccionario observado de resultado academico A/E/I/R/NULL."
        uso_prohibido = "No inferir significado de A/E/I/R por nombre o intuicion."
        estado = "GOBERNADA"
        justificacion = "Diccionario explicito detectado en PROMEDIOS y corroborado por paso 50."

    if normalized in {"ANO", "PERIODO", "CODRAMO", "RAMOEQUIV"}:
        role = "periodo" if normalized in {"ANO", "PERIODO"} else "codigo"
        uso_permitido = "Columna fuente para filtrar año 2025, periodo y ramos distintos."
        uso_prohibido = "No mezclar año proceso 2026 con año academico 2025."
        estado = "GOBERNADA"
        justificacion = "Columna critica para formulas 16-19; requiere filtros documentados."

    if normalized in {"CURSO_1ER_SEM", "CURSO_2DO_SEM", "UNIDADES_CURSADAS", "UNIDADES_APROBADAS"}:
        role = "metrica"
        uso_permitido = "Campo oficial 5809 anual 2025; solo recalcular si reglas y bloqueos cierran."
        uso_prohibido = "No recalcular definitivo en esta etapa de gobernanza."
        estado = "PENDIENTE"
        justificacion = "Regla documentada, pero existen bloqueos actuales post paso 50."

    if normalized in {"UNID_CURSADAS_TOTAL", "UNID_APROBADAS_TOTAL"}:
        role = "metrica"
        uso_permitido = "Campo acumulado hasta cierre 2025; requiere fuente historica y regla cerrada."
        uso_prohibido = "No asumir igualdad con anual ni usar planes como numerador."
        estado = "PENDIENTE"
        justificacion = "Acumulado 20-21 no esta cerrado."

    if normalized == "VIGENCIA":
        role = "estado"
        uso_permitido = "Mantener/eliminar registro segun codigos oficiales 0/1."
        uso_prohibido = "No usar estado academico para forzar vigencia o unidades cero."
        estado = "GOBERNADA"
        justificacion = "Campo oficial con codigos 0/1 del instructivo."

    if normalized in {"ESTADOACADEMICO", "ESTADO_ACADEMICO", "SITUACION", "CLASIFICACION_IDENTIDAD"}:
        role = "estado"
        uso_permitido = "Auditoria o clasificacion observada."
        uso_prohibido = "No usar para forzar avance academico ni reemplazar ESTADO/DESCRIPCION_ESTADO."
        estado = "PENDIENTE"
        justificacion = "Catalogo observado; requiere validacion funcional antes de uso operativo."

    if normalized in {
        "CODCARR",
        "CODCARPR",
        "CODCARR_ESPERADO_SEGUN_CARRERAS",
        "CODCARR_ESPERADO_EN_DATOS",
        "CODPESTUD",
        "PLAN_DE_ESTUDIO",
        "NOMBRE_CARRERA",
        "NOMBRE_SEDE",
        "JORNADA",
        "VERSION",
        "SEDE",
    }:
        role = "codigo"
        uso_permitido = "Referencia de programa, plan, sede o jornada para validacion."
        uso_prohibido = "No usar para cambiar regla oficial ni para inventar llaves no conciliadas."
        estado = "GOBERNADA" if normalized in {"CODIGO_UNICO", "PLAN_ESTUDIOS"} else "PENDIENTE"
        justificacion = "Catalogo o atributo de programa; uso depende del cruce gobernado."

    if normalized in {"FECHA_NACIMIENTO", "FECHAMATRICULA", "FECHAFIRMA"}:
        role = "fecha"

    return role, uso_permitido, uso_prohibido, nivel, estado, justificacion


def inventory_txt(
    source_id: str,
    path: Path,
    hoja_id: str,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    lines = path.read_text(encoding="utf-8-sig", errors="replace").splitlines()
    stats = make_stats(["LINEA_TEXTO"])
    for line in lines:
        update_stats(stats, [line])
    sheet_class = classify_sheet(source_id, "TXT")
    hoja = {
        "ID_FUENTE": source_id,
        "ID_HOJA": hoja_id,
        "Archivo": path.name,
        "Hoja": "TXT",
        "Filas": len(lines),
        "Columnas": 1,
        **sheet_class,
    }
    columnas = []
    for item in finish_stats(stats):
        role, permit, prohibit, nivel, estado, just = classify_column(
            source_id, "TXT", item["original"], item["normalizado"], item["posicion"]
        )
        columnas.append(
            {
                "ID_FUENTE": source_id,
                "ID_HOJA": hoja_id,
                "ID_COLUMNA": f"{hoja_id}_C{item['posicion']:03d}",
                "Archivo": path.name,
                "Hoja": "TXT",
                "Nombre columna original": item["original"],
                "Nombre normalizado": item["normalizado"],
                "Posicion": item["posicion"],
                "Tipo observado": item["tipo_observado"],
                "No vacios": item["no_vacios"],
                "Distintos": item["distintos"],
                "Muestra valores": item["muestra"],
                "Posible rol": role,
                "Uso permitido": permit,
                "Uso prohibido": prohibit,
                "Nivel respaldo": nivel,
                "Estado gobernanza": estado,
                "Justificacion": just,
            }
        )
    return hoja, columnas


def inventory_csv(
    source_id: str,
    path: Path,
    hoja_id: str,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]], List[str]]:
    encoding, sep = detect_csv(path)
    rows_count = 0
    with path.open("r", encoding=encoding, errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=sep)
        headers = next(reader)
        stats = make_stats(headers)
        for row in reader:
            rows_count += 1
            update_stats(stats, row)
    sheet_class = classify_sheet(source_id, "CSV")
    sheet_class["Observacion"] += f" Lectura tecnica: encoding={encoding}; separador={repr(sep)}."
    hoja = {
        "ID_FUENTE": source_id,
        "ID_HOJA": hoja_id,
        "Archivo": path.name,
        "Hoja": "CSV",
        "Filas": rows_count,
        "Columnas": len(headers),
        **sheet_class,
    }
    columnas = []
    for item in finish_stats(stats):
        role, permit, prohibit, nivel, estado, just = classify_column(
            source_id, "CSV", item["original"], item["normalizado"], item["posicion"]
        )
        columnas.append(
            {
                "ID_FUENTE": source_id,
                "ID_HOJA": hoja_id,
                "ID_COLUMNA": f"{hoja_id}_C{item['posicion']:03d}",
                "Archivo": path.name,
                "Hoja": "CSV",
                "Nombre columna original": item["original"],
                "Nombre normalizado": item["normalizado"],
                "Posicion": item["posicion"],
                "Tipo observado": item["tipo_observado"],
                "No vacios": item["no_vacios"],
                "Distintos": item["distintos"],
                "Muestra valores": item["muestra"],
                "Posible rol": role,
                "Uso permitido": permit,
                "Uso prohibido": prohibit,
                "Nivel respaldo": nivel,
                "Estado gobernanza": estado,
                "Justificacion": just,
            }
        )
    return hoja, columnas, [display_value(h) for h in headers]


def inventory_xlsx(
    source_id: str,
    path: Path,
    sheet_counter_start: int,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], int]:
    hojas: List[Dict[str, Any]] = []
    columnas: List[Dict[str, Any]] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    counter = sheet_counter_start
    try:
        for ws in wb.worksheets:
            hoja_id = f"{source_id}_H{counter:02d}"
            counter += 1
            rows_iter = ws.iter_rows(values_only=True)
            try:
                raw_headers = list(next(rows_iter))
            except StopIteration:
                raw_headers = []
            max_columns = max(ws.max_column or 0, len(raw_headers))
            if len(raw_headers) < max_columns:
                raw_headers += [None] * (max_columns - len(raw_headers))
            stats = make_stats(raw_headers)
            data_rows = 0
            for row in rows_iter:
                data_rows += 1
                update_stats(stats, list(row))
            sheet_class = classify_sheet(source_id, ws.title)
            hojas.append(
                {
                    "ID_FUENTE": source_id,
                    "ID_HOJA": hoja_id,
                    "Archivo": path.name,
                    "Hoja": ws.title,
                    "Filas": data_rows,
                    "Columnas": max_columns,
                    **sheet_class,
                }
            )
            for item in finish_stats(stats):
                role, permit, prohibit, nivel, estado, just = classify_column(
                    source_id, ws.title, item["original"], item["normalizado"], item["posicion"]
                )
                columnas.append(
                    {
                        "ID_FUENTE": source_id,
                        "ID_HOJA": hoja_id,
                        "ID_COLUMNA": f"{hoja_id}_C{item['posicion']:03d}",
                        "Archivo": path.name,
                        "Hoja": ws.title,
                        "Nombre columna original": item["original"],
                        "Nombre normalizado": item["normalizado"],
                        "Posicion": item["posicion"],
                        "Tipo observado": item["tipo_observado"],
                        "No vacios": item["no_vacios"],
                        "Distintos": item["distintos"],
                        "Muestra valores": item["muestra"],
                        "Posible rol": role,
                        "Uso permitido": permit,
                        "Uso prohibido": prohibit,
                        "Nivel respaldo": nivel,
                        "Estado gobernanza": estado,
                        "Justificacion": just,
                    }
                )
    finally:
        wb.close()
    return hojas, columnas, counter


def build_inventories() -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, Dict[str, str], List[str]]:
    fecha_lectura = now_iso()
    source_rows = []
    sheet_rows = []
    column_rows = []
    hashes = {}
    headers_5809: List[str] = []
    for source_id, meta in PATHS.items():
        path = meta["ruta"]
        if not path.exists():
            raise FileNotFoundError(path)
        digest = sha256_file(path)
        hashes[source_id] = digest
        source_rows.append(
            {
                "ID_FUENTE": source_id,
                "Ruta": str(path),
                "Nombre archivo": path.name,
                "Tipo de archivo": path.suffix.lower().lstrip(".") or "txt",
                "Clasificacion": meta["clasificacion"],
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Origen": meta["origen"],
                "Nivel de respaldo": meta["nivel"],
                "Puede usarse para regla oficial": meta["regla"],
                "Puede usarse para dato observado": meta["dato"],
                "Puede usarse para implementacion tecnica": meta["implementacion"],
                "Puede modificar regla oficial": "NO",
                "Estado": meta["estado"],
                "Hash SHA256": digest,
                "Fecha lectura": fecha_lectura,
                "Observacion": meta["observacion"],
            }
        )
        ext = path.suffix.lower()
        if ext == ".txt":
            hoja, cols = inventory_txt(source_id, path, f"{source_id}_H01")
            sheet_rows.append(hoja)
            column_rows.extend(cols)
        elif ext == ".csv":
            hoja, cols, headers = inventory_csv(source_id, path, f"{source_id}_H01")
            sheet_rows.append(hoja)
            column_rows.extend(cols)
            if source_id == "F02_PRECARGA_5809":
                headers_5809 = headers
        elif ext in {".xlsx", ".xlsm"}:
            hojas, cols, _ = inventory_xlsx(source_id, path, 1)
            sheet_rows.extend(hojas)
            column_rows.extend(cols)
        else:
            raise ValueError(f"Tipo no soportado: {path}")
    return (
        pd.DataFrame(source_rows),
        pd.DataFrame(sheet_rows),
        pd.DataFrame(column_rows),
        hashes,
        headers_5809,
    )


def read_xlsx_rows(path: Path, sheet_name: str) -> List[Dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: List[Dict[str, Any]] = []
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        headers = [display_value(h) for h in next(iterator)]
        for row in iterator:
            rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    finally:
        wb.close()
    return rows


def unique_pairs_xlsx(path: Path, sheet_name: str, code_col: str, desc_col: str, limit: int | None = None) -> List[Tuple[str, str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    pairs = set()
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        headers = [normalize_text(h) for h in next(iterator)]
        code_idx = headers.index(normalize_text(code_col))
        desc_idx = headers.index(normalize_text(desc_col))
        for row in iterator:
            code = row[code_idx] if code_idx < len(row) else None
            desc = row[desc_idx] if desc_idx < len(row) else None
            code_text = display_value(code)
            desc_text = display_value(desc)
            if code_text.upper() == "NULL" or not code_text:
                code_text = "NULL / blanco"
            if desc_text.upper() == "NULL" or not desc_text:
                desc_text = "(en blanco)"
            pairs.add((code_text, desc_text))
            if limit and len(pairs) >= limit:
                break
    finally:
        wb.close()
    return sorted(pairs, key=lambda x: (x[0], x[1]))


def unique_values_xlsx(path: Path, sheet_name: str, col: str, limit: int = 100) -> List[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    values = set()
    try:
        ws = wb[sheet_name]
        iterator = ws.iter_rows(values_only=True)
        headers = [normalize_text(h) for h in next(iterator)]
        idx = headers.index(normalize_text(col))
        for row in iterator:
            val = row[idx] if idx < len(row) else None
            if not is_empty(val):
                values.add(display_value(val))
            if len(values) >= limit:
                break
    finally:
        wb.close()
    return sorted(values)


def unique_pairs_csv(path: Path, code_col: str, desc_col: str | None = None) -> List[Tuple[str, str]]:
    encoding, sep = detect_csv(path)
    pairs = set()
    with path.open("r", encoding=encoding, errors="replace", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=sep)
        for row in reader:
            code = display_value(row.get(code_col))
            if not code:
                continue
            if desc_col:
                desc = display_value(row.get(desc_col))
            else:
                desc = code
            pairs.add((code, desc))
    return sorted(pairs, key=lambda x: (x[0], x[1]))


def build_dictionaries() -> Tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    prom_path = PATHS["F04_PROMEDIOS_ACTUALIZADO"]["ruta"]
    recalc_path = PATHS["F07_RECALCULO_50_EVIDENCIA"]["ruta"]
    carreras_path = PATHS["F03_PRECARGA_5810"]["ruta"]
    planes_path = PATHS["F06_PLANES_ESTUDIO"]["ruta"]
    mapeo_path = PATHS["F05_MAPEO_IDENTIDAD_CODCLI"]["ruta"]

    for code, desc in unique_pairs_xlsx(prom_path, "Hoja1", "ESTADO", "DESCRIPCION_ESTADO"):
        treatment = {
            "A": "APROBADA",
            "E": "APROBADA_CONVALIDADA_CON_VALIDACION_FUNCIONAL",
            "I": "APROBADA_HOMOLOGADA_CON_VALIDACION_FUNCIONAL",
            "R": "NO_APROBADA",
            "NULL / blanco": "REVISION_NO_AUTOMATICO",
        }.get(code, "REVISION")
        rows.append(
            {
                "ID_DICCIONARIO": "DIC_ESTADO_PROMEDIOS",
                "Fuente": str(prom_path),
                "Hoja": "Hoja1",
                "Campo codigo": "ESTADO",
                "Campo descripcion": "DESCRIPCION_ESTADO",
                "Codigo": code,
                "Descripcion": desc,
                "Tratamiento tecnico propuesto": treatment,
                "Nivel respaldo": "B",
                "Puede usarse en calculo": "SI_CON_VALIDACION" if code in {"E", "I"} else ("NO" if code == "NULL / blanco" else "SI"),
                "Requiere validacion funcional": "SI" if code in {"E", "I", "NULL / blanco"} else "NO",
                "Observacion": "Diccionario explicito detectado antes de inferir estados.",
            }
        )

    for row in read_xlsx_rows(recalc_path, "01_DICCIONARIO_ESTADO"):
        rows.append(
            {
                "ID_DICCIONARIO": "DIC_ESTADO_RECALCULO_50",
                "Fuente": str(recalc_path),
                "Hoja": "01_DICCIONARIO_ESTADO",
                "Campo codigo": "ESTADO",
                "Campo descripcion": "DESCRIPCION_ESTADO",
                "Codigo": display_value(row.get("ESTADO")) or "NULL / blanco",
                "Descripcion": display_value(row.get("DESCRIPCION_ESTADO")),
                "Tratamiento tecnico propuesto": display_value(row.get("TRATAMIENTO_19_APROBADAS")),
                "Nivel respaldo": "C",
                "Puede usarse en calculo": "NO_COMO_REGLA; SI_COMO_EVIDENCIA",
                "Requiere validacion funcional": "SI",
                "Observacion": "Evidencia tecnica no normativa del paso 50.",
            }
        )

    for value in unique_values_xlsx(prom_path, "Hoja1", "ESTADO_ACADEMICO", 100):
        rows.append(
            {
                "ID_DICCIONARIO": "CAT_ESTADO_ACADEMICO_PROMEDIOS",
                "Fuente": str(prom_path),
                "Hoja": "Hoja1",
                "Campo codigo": "ESTADO_ACADEMICO",
                "Campo descripcion": "ESTADO_ACADEMICO",
                "Codigo": value,
                "Descripcion": value,
                "Tratamiento tecnico propuesto": "Solo auditoria; no forzar unidades ni vigencia.",
                "Nivel respaldo": "B",
                "Puede usarse en calculo": "NO",
                "Requiere validacion funcional": "SI",
                "Observacion": "Catalogo observado de estado academico; no equivale a avance curricular.",
            }
        )

    for code, desc in [("0", "Eliminar Registro"), ("1", "Mantener Registro")]:
        rows.append(
            {
                "ID_DICCIONARIO": "DIC_VIGENCIA_OFICIAL_5809_5810",
                "Fuente": str(PATHS["F01_INSTRUCTIVO_OFICIAL"]["ruta"]),
                "Hoja": "TXT",
                "Campo codigo": "VIGENCIA",
                "Campo descripcion": "VIGENCIA",
                "Codigo": code,
                "Descripcion": desc,
                "Tratamiento tecnico propuesto": "Validar codigo oficial; no derivar desde estado academico.",
                "Nivel respaldo": "A",
                "Puede usarse en calculo": "SI",
                "Requiere validacion funcional": "NO",
                "Observacion": "Catalogo oficial declarado en instructivo.",
            }
        )

    for code, desc in unique_pairs_csv(carreras_path, "CODIGO_UNICO", "NOMBRE_CARRERA"):
        rows.append(
            {
                "ID_DICCIONARIO": "CAT_CARRERAS_5810_CODIGO_UNICO",
                "Fuente": str(carreras_path),
                "Hoja": "CSV",
                "Campo codigo": "CODIGO_UNICO",
                "Campo descripcion": "NOMBRE_CARRERA",
                "Codigo": code,
                "Descripcion": desc,
                "Tratamiento tecnico propuesto": "Referencia de programa; cruzar con PLAN_ESTUDIOS.",
                "Nivel respaldo": "A",
                "Puede usarse en calculo": "SI_COMO_VALIDACION",
                "Requiere validacion funcional": "NO",
                "Observacion": "Catalogo observado en precarga de carreras.",
            }
        )

    for field in ["PLAN_ESTUDIOS", "JORNADA", "VERSION", "NOMBRE_SEDE"]:
        for code, desc in unique_pairs_csv(carreras_path, field, field):
            rows.append(
                {
                    "ID_DICCIONARIO": f"CAT_5810_{field}",
                    "Fuente": str(carreras_path),
                    "Hoja": "CSV",
                    "Campo codigo": field,
                    "Campo descripcion": field,
                    "Codigo": code,
                    "Descripcion": desc,
                    "Tratamiento tecnico propuesto": "Catalogo de referencia de carreras 5810.",
                    "Nivel respaldo": "A",
                    "Puede usarse en calculo": "SI_COMO_VALIDACION",
                    "Requiere validacion funcional": "NO",
                    "Observacion": "Catalogo observado en precarga.",
                }
            )

    for code, desc in unique_pairs_xlsx(planes_path, "BBDD Bruta", "CODPESTUD", "NOMPESTUD"):
        rows.append(
            {
                "ID_DICCIONARIO": "CAT_PLANES_ESTUDIO_CODPESTUD",
                "Fuente": str(planes_path),
                "Hoja": "BBDD Bruta",
                "Campo codigo": "CODPESTUD",
                "Campo descripcion": "NOMPESTUD",
                "Codigo": code,
                "Descripcion": desc,
                "Tratamiento tecnico propuesto": "Referencia de plan; denominador/validacion, no numerador automatico.",
                "Nivel respaldo": "B",
                "Puede usarse en calculo": "SI_COMO_REFERENCIA",
                "Requiere validacion funcional": "SI",
                "Observacion": "Catalogo institucional de planes.",
            }
        )

    for value in unique_values_xlsx(planes_path, "BBDD Bruta", "Sede", 100):
        rows.append(
            {
                "ID_DICCIONARIO": "CAT_SEDE_PLANES_ESTUDIO",
                "Fuente": str(planes_path),
                "Hoja": "BBDD Bruta",
                "Campo codigo": "Sede",
                "Campo descripcion": "Sede",
                "Codigo": value,
                "Descripcion": value,
                "Tratamiento tecnico propuesto": "Referencia de sede; no altera llave 5809.",
                "Nivel respaldo": "B",
                "Puede usarse en calculo": "SI_COMO_REFERENCIA",
                "Requiere validacion funcional": "SI",
                "Observacion": "Catalogo observado de sede en planes.",
            }
        )

    for row in read_xlsx_rows(mapeo_path, "RESUMEN"):
        rows.append(
            {
                "ID_DICCIONARIO": "CAT_CLASIFICACION_IDENTIDAD_MAPEO",
                "Fuente": str(mapeo_path),
                "Hoja": "RESUMEN",
                "Campo codigo": "CLASIFICACION",
                "Campo descripcion": "CLASIFICACION",
                "Codigo": display_value(row.get("CLASIFICACION")),
                "Descripcion": display_value(row.get("CLASIFICACION")),
                "Tratamiento tecnico propuesto": f"Auditoria de identidad; casos={display_value(row.get('CASOS'))}.",
                "Nivel respaldo": "B",
                "Puede usarse en calculo": "SI_COMO_CONTROL",
                "Requiere validacion funcional": "SI",
                "Observacion": "Catalogo de clasificacion de conciliacion.",
            }
        )

    df = pd.DataFrame(rows)
    df_estado = pd.DataFrame(DICT_ESTADO_PROMEDIOS)
    return df, df_estado


def build_keys() -> pd.DataFrame:
    rows = [
        {
            "ID_LLAVE": "K01_5809_OFICIAL",
            "Nombre llave": "Llave 5809 oficial",
            "Columnas usadas": "posición/fila, CODIGO_UNICO, PLAN_ESTUDIOS, NUM_DOCUMENTO, DV",
            "Archivos involucrados": "5809 precarga; instructivo; mapeo identidad",
            "Proposito": "Mantener trazabilidad de cada fila oficial y programa informado.",
            "Cardinalidad esperada": "Una fila oficial 5809 por posicion; CODIGO_UNICO+PLAN puede repetir por estudiantes.",
            "Duplicados permitidos": "SI para programa; NO para ID_FILA_5809.",
            "Riesgo": "Modificar atributos no modificables o resolver por nombres.",
            "Estado gobernanza": "GOBERNADA",
            "Evidencia": "F02_PRECARGA_5809; F05_MAPEO_IDENTIDAD_CODCLI.",
        },
        {
            "ID_LLAVE": "K02_RUT_NORMALIZADO",
            "Nombre llave": "Llave tecnica RUT normalizado",
            "Columnas usadas": "NUM_DOCUMENTO + DV / RUT + DIG",
            "Archivos involucrados": "5809; PROMEDIOS; MAPEO_IDENTIDAD_CODCLI",
            "Proposito": "Conciliar identidad documental entre precarga y fuentes institucionales.",
            "Cardinalidad esperada": "Puede haber multiples CODCLI por RUT.",
            "Duplicados permitidos": "SI, con auditoria de multiplicidad.",
            "Riesgo": "Sumar todos los CODCLI de un RUT sin filtrar por programa.",
            "Estado gobernanza": "GOBERNADA_CON_RESTRICCION",
            "Evidencia": "Campos DOC_PRECARGA_FULL, RUT_INSTITUCIONAL, CODCLI_LISTA.",
        },
        {
            "ID_LLAVE": "K03_CODCLI_ACADEMICO",
            "Nombre llave": "Llave tecnica CODCLI",
            "Columnas usadas": "CODCLI / CODCLI_LISTA",
            "Archivos involucrados": "PROMEDIOS Hoja1; MAPEO MAPEO",
            "Proposito": "Vincular fila 5809 con registros academicos.",
            "Cardinalidad esperada": "Uno o mas CODCLI candidatos; debe filtrarse por programa.",
            "Duplicados permitidos": "SI en lista; NO como seleccion final no auditada.",
            "Riesgo": "Usar N_CODCLI como CODCLI academico.",
            "Estado gobernanza": "GOBERNADA",
            "Evidencia": "CODCLI_LISTA documentado como columna correcta; N_CODCLI bloqueado.",
        },
        {
            "ID_LLAVE": "K04_PROGRAMA",
            "Nombre llave": "Llave programa",
            "Columnas usadas": "RUT + CODIGO_UNICO + PLAN_ESTUDIOS",
            "Archivos involucrados": "5809; 5810; MAPEO; PROMEDIOS matriz/planes",
            "Proposito": "Evitar cruces de CODCLI de otra carrera del mismo RUT.",
            "Cardinalidad esperada": "Una relacion gobernada por fila 5809-programa.",
            "Duplicados permitidos": "NO sin decision funcional.",
            "Riesgo": "Cruzar por RUT solamente.",
            "Estado gobernanza": "GOBERNADA_CON_RESTRICCION",
            "Evidencia": "Instructivo exige carrera informada en matricula 2025.",
        },
        {
            "ID_LLAVE": "K05_DETALLE_ACADEMICO",
            "Nombre llave": "Llave de detalle academico",
            "Columnas usadas": "CODCLI + ANO + PERIODO + CODRAMO",
            "Archivos involucrados": "PROMEDIOS Hoja1",
            "Proposito": "Identificar actividad academica por año/periodo/ramo.",
            "Cardinalidad esperada": "Puede haber mas de un registro si existen secciones/equivalencias; contar CODRAMO unico requiere regla.",
            "Duplicados permitidos": "SI en raw; NO en conteo sin deduplicacion gobernada.",
            "Riesgo": "Contar intentos/secciones como unidades distintas.",
            "Estado gobernanza": "GOBERNADA",
            "Evidencia": "Columnas CODCLI, ANO, PERIODO, CODRAMO detectadas en PROMEDIOS Hoja1.",
        },
        {
            "ID_LLAVE": "K06_UNIDADES_CURSADAS",
            "Nombre llave": "Llave para contar unidades cursadas",
            "Columnas usadas": "CODCLI filtrado + CODRAMO + ANO=2025",
            "Archivos involucrados": "MAPEO; PROMEDIOS Hoja1",
            "Proposito": "Contar ramos/unidades cursadas del año academico 2025.",
            "Cardinalidad esperada": "Un CODRAMO unico por unidad contada.",
            "Duplicados permitidos": "NO despues de deduplicar CODRAMO.",
            "Riesgo": "Incluir ramos de otra carrera o periodos no gobernados.",
            "Estado gobernanza": "GOBERNADA_CON_BLOQUEOS",
            "Evidencia": "Paso 50 identifica 122 sin registros 2025 y 71 diferencias 16-18 o multiples.",
        },
        {
            "ID_LLAVE": "K07_UNIDADES_APROBADAS",
            "Nombre llave": "Llave para contar unidades aprobadas",
            "Columnas usadas": "CODCLI filtrado + CODRAMO + estado aprobatorio",
            "Archivos involucrados": "PROMEDIOS Hoja1; diccionario ESTADO/DESCRIPCION_ESTADO",
            "Proposito": "Contar unidades aprobadas segun diccionario y regla funcional.",
            "Cardinalidad esperada": "Un CODRAMO aprobado unico por unidad contada.",
            "Duplicados permitidos": "NO despues de deduplicar y resolver E/I/NULL.",
            "Riesgo": "Aplicar convalidacion/homologacion anual sin validacion funcional.",
            "Estado gobernanza": "GOBERNADA_CON_BLOQUEOS",
            "Evidencia": "Diccionario A/E/I/R/NULL detectado; 223 diferencias solo aprobadas 19.",
        },
    ]
    return pd.DataFrame(rows)


def official_rule_for_col(col: str, pos: int) -> Dict[str, str]:
    rules = {
        "CURSO_1ER_SEM": (
            "Indica si cursó actividades academicas exigidas durante el primer semestre 2025.",
            "PROMEDIOS Hoja1; CODCLI_LISTA; ANO; PERIODO; CODRAMO.",
            "SI/NO por presencia de actividad gobernada en PERIODO=1; PERIODO=3 queda en revision.",
            "Valores SI/NO; si SI, UNIDADES_CURSADAS > 0.",
            "GOBERNADA_CON_BLOQUEOS",
            "Resolver pendientes paso 50 y periodos no 1/2.",
        ),
        "CURSO_2DO_SEM": (
            "Indica si cursó actividades academicas exigidas durante el segundo semestre 2025.",
            "PROMEDIOS Hoja1; CODCLI_LISTA; ANO; PERIODO; CODRAMO.",
            "SI/NO por presencia de actividad gobernada en PERIODO=2; PERIODO=3 queda en revision.",
            "Valores SI/NO; si SI, UNIDADES_CURSADAS > 0.",
            "GOBERNADA_CON_BLOQUEOS",
            "Resolver pendientes paso 50 y periodos no 1/2.",
        ),
        "UNIDADES_CURSADAS": (
            "Número de unidades de medida cursadas efectivamente durante año academico 2025.",
            "PROMEDIOS Hoja1; CODCLI_LISTA filtrado; CODRAMO; ANO=2025.",
            "Conteo de CODRAMO unico con resultado gobernado, separado de NULL/blanco.",
            "No mayor que total; coherencia con curso semestre.",
            "GOBERNADA_CON_BLOQUEOS",
            "Resolver 71 diferencias 16-18/multiple y 122 sin registros 2025.",
        ),
        "UNIDADES_APROBADAS": (
            "Número de unidades aprobadas efectivamente durante año academico 2025.",
            "PROMEDIOS Hoja1; ESTADO/DESCRIPCION_ESTADO; CODCLI_LISTA.",
            "Conteo de CODRAMO unico con A/E/I segun diccionario, con E/I sujetos a validacion funcional anual.",
            "No mayor que cursadas ni que total aprobadas.",
            "PENDIENTE_VALIDACION_FUNCIONAL",
            "Resolver tratamiento anual de E/I y NULL/blanco; 223 diferencias solo aprobadas 19.",
        ),
        "UNID_CURSADAS_TOTAL": (
            "Unidades cursadas acumuladas desde ingreso hasta cierre academico 2025.",
            "Fuente historica academica y planes; PROMEDIOS no queda cerrado sin validacion.",
            "No calcular en esta etapa. Separar anual 16-19 de acumulado 20-21.",
            "Debe ser >= anual cuando corresponda; coherencia con plan.",
            "PENDIENTE",
            "Definir fuente historica completa, homologaciones, convalidaciones y deduplicacion acumulada.",
        ),
        "UNID_APROBADAS_TOTAL": (
            "Unidades aprobadas acumuladas desde ingreso hasta cierre academico 2025.",
            "Fuente historica academica, convalidaciones/homologaciones, planes como limite de referencia.",
            "No calcular en esta etapa. Planes son denominador/referencia, no numerador.",
            "No puede exceder cursadas total; tolerancia oficial contra total plan.",
            "PENDIENTE",
            "Cerrar regla acumulada y limite por plan.",
        ),
        "VIGENCIA": (
            "Mantener o eliminar registro cargado: 0 eliminar, 1 mantener.",
            "Precarga 5809 e instructivo.",
            "Validar codigo; no derivar automaticamente desde estado academico.",
            "Valores 0/1.",
            "GOBERNADA",
            "Sin pendiente para lectura; no generar carga.",
        ),
    }
    if col in rules:
        regla, fuente, transf, validacion, estado, pendiente = rules[col]
        return {
            "Regla oficial si existe": regla,
            "Fuente de dato": fuente,
            "Transformacion": transf,
            "Validacion": validacion,
            "Estado actual": estado,
            "Pendiente": pendiente,
        }
    if pos <= 15:
        return {
            "Regla oficial si existe": "Dato precargado/no modificable segun instructivo o base oficial.",
            "Fuente de dato": "Precarga oficial 5809.",
            "Transformacion": "Ninguna; conservar valor y orden.",
            "Validacion": "Comparar contra precarga original; no editar atributos no modificables.",
            "Estado actual": "GOBERNADA_RAW",
            "Pendiente": "Sin recalculo; mantener trazabilidad.",
        }
    return {
        "Regla oficial si existe": "Regla no documentada en esta matriz.",
        "Fuente de dato": "Pendiente.",
        "Transformacion": "Pendiente.",
        "Validacion": "Pendiente.",
        "Estado actual": "PENDIENTE",
        "Pendiente": "Gobernar antes de uso.",
    }


def build_mapeo_5809(headers_5809: List[str]) -> pd.DataFrame:
    columns = headers_5809 or OFFICIAL_5809_COLUMNS
    rows = []
    for pos, col in enumerate(columns, start=1):
        rule = official_rule_for_col(col, pos)
        rows.append(
            {
                "Numero columna": pos,
                "Nombre columna": col,
                "Fuente oficial": "Instructivo oficial Avance Curricular SIES 2026 + Precarga 5809",
                **rule,
            }
        )
    return pd.DataFrame(rows)


def build_governance_16_19() -> pd.DataFrame:
    common = {
        "Anio de referencia": 2025,
        "Fuente primaria": "PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx",
        "Hoja fuente": "Hoja1",
        "Llave usada": "CODCLI_LISTA desde MAPEO, filtrado por programa; luego CODCLI + ANO + PERIODO + CODRAMO.",
        "Diccionario usado": "ESTADO / DESCRIPCION_ESTADO de PROMEDIOS; hoja 05 de esta gobernanza.",
        "Exclusiones": "No usar N_CODCLI; no sumar CODCLI de mismo RUT sin programa; no mezclar 2026; no usar NULL/blanco automatico.",
        "Validaciones": "Hash fuente, hoja correcta, columnas criticas, llave, cardinalidad, ANO=2025, PERIODO, CODCLI_LISTA, diccionario estado.",
    }
    rows = [
        {
            "Columna 5809": "16 CURSO_1ER_SEM",
            "Que pregunta responde": "Si el/la estudiante cursó actividades academicas durante primer semestre 2025.",
            "Columnas fuente usadas": "CODCLI, ANO, PERIODO, CODRAMO, ESTADO, DESCRIPCION_ESTADO.",
            "Formula": "SI si existe al menos un CODRAMO gobernado con ANO=2025 y PERIODO=1 para CODCLI filtrado; NO si no existe y no hay bloqueo.",
            "Inclusiones": "Registros 2025 del programa informado; periodo 1 gobernado.",
            "Tratamiento de convalidacion": "No determina por si sola presencia semestral sin validar periodo/programa.",
            "Tratamiento de homologacion": "No determina por si sola presencia semestral sin validar periodo/programa.",
            "Tratamiento de reprobacion": "Cuenta como presencia/cursada si pertenece al periodo/programa.",
            "Tratamiento de NULL/blanco": "Revision; no aplicar automatico.",
            "Pendientes": "PERIODO distinto de 1/2 y filas bloqueadas paso 50.",
            "Estado": "GOBERNADA_CON_BLOQUEOS",
            **common,
        },
        {
            "Columna 5809": "17 CURSO_2DO_SEM",
            "Que pregunta responde": "Si el/la estudiante cursó actividades academicas durante segundo semestre 2025.",
            "Columnas fuente usadas": "CODCLI, ANO, PERIODO, CODRAMO, ESTADO, DESCRIPCION_ESTADO.",
            "Formula": "SI si existe al menos un CODRAMO gobernado con ANO=2025 y PERIODO=2 para CODCLI filtrado; NO si no existe y no hay bloqueo.",
            "Inclusiones": "Registros 2025 del programa informado; periodo 2 gobernado.",
            "Tratamiento de convalidacion": "No determina por si sola presencia semestral sin validar periodo/programa.",
            "Tratamiento de homologacion": "No determina por si sola presencia semestral sin validar periodo/programa.",
            "Tratamiento de reprobacion": "Cuenta como presencia/cursada si pertenece al periodo/programa.",
            "Tratamiento de NULL/blanco": "Revision; no aplicar automatico.",
            "Pendientes": "PERIODO distinto de 1/2 y filas bloqueadas paso 50.",
            "Estado": "GOBERNADA_CON_BLOQUEOS",
            **common,
        },
        {
            "Columna 5809": "18 UNIDADES_CURSADAS",
            "Que pregunta responde": "Cuantas unidades de medida cursó efectivamente en el año academico 2025.",
            "Columnas fuente usadas": "CODCLI, ANO, CODRAMO, ESTADO, DESCRIPCION_ESTADO, CODCARR/PLAN si aplica.",
            "Formula": "Contar CODRAMO unico con ANO=2025 del CODCLI filtrado por programa, separando NULL/blanco.",
            "Inclusiones": "A, R y otros resultados gobernados como actividad cursada; E/I solo si corresponden a actividad aplicable al campo.",
            "Tratamiento de convalidacion": "El anual no debe incluir validaciones de otras carreras sin decision funcional.",
            "Tratamiento de homologacion": "El anual no debe incluir homologaciones sin decision funcional.",
            "Tratamiento de reprobacion": "Cuenta como cursada; no cuenta como aprobada.",
            "Tratamiento de NULL/blanco": "Revision; no contar automatico.",
            "Pendientes": "71 diferencias 16-18/multiple y 122 sin registros 2025.",
            "Estado": "GOBERNADA_CON_BLOQUEOS",
            **common,
        },
        {
            "Columna 5809": "19 UNIDADES_APROBADAS",
            "Que pregunta responde": "Cuantas unidades de medida aprobó durante el año academico 2025.",
            "Columnas fuente usadas": "CODCLI, ANO, CODRAMO, ESTADO, DESCRIPCION_ESTADO.",
            "Formula": "Contar CODRAMO unico con estado aprobatorio segun diccionario: A; E/I solo con validacion funcional; R no; NULL revision.",
            "Inclusiones": "A como aprobado; E/I como aprobatorio observado sujeto a regla por campo.",
            "Tratamiento de convalidacion": "E=CONVALIDACION requiere confirmar compatibilidad con exclusion anual del instructivo.",
            "Tratamiento de homologacion": "I=HOMOLOGADO requiere confirmar compatibilidad con exclusion anual del instructivo.",
            "Tratamiento de reprobacion": "R no suma como aprobada.",
            "Tratamiento de NULL/blanco": "Revision; no contar automatico.",
            "Pendientes": "223 diferencias solo aprobadas 19; decision funcional E/I/NULL.",
            "Estado": "PENDIENTE_VALIDACION_FUNCIONAL",
            **common,
        },
    ]
    return pd.DataFrame(rows)


def build_governance_20_21() -> pd.DataFrame:
    rows = [
        {
            "Tema": "Definicion acumulado",
            "Detalle": "Acumulado desde ingreso del estudiante a la carrera hasta cierre academico 2025.",
            "Fuente requerida": "Historial academico completo por CODCLI/programa, con equivalencias, homologaciones y convalidaciones.",
            "PROMEDIOS permite calcularlo": "NO_CERRADO",
            "Planes de estudio": "Denominador/limite/referencia, no numerador automatico.",
            "Que no se puede asumir": "No asumir que anual 2025 equivale a total; no usar total de plan como aprobado; no mezclar carreras.",
            "Bloqueos actuales": "20-21 acumulado pendiente/no cerrado.",
            "Estado": "PENDIENTE",
        },
        {
            "Tema": "UNID_CURSADAS_TOTAL",
            "Detalle": "Unidades cursadas acumuladas, incluyendo reglas de validacion/homologacion cuando corresponda.",
            "Fuente requerida": "Detalle historico por CODCLI y plan desde ingreso a cierre 2025.",
            "PROMEDIOS permite calcularlo": "PENDIENTE_DE_VALIDAR_COBERTURA_HISTORICA",
            "Planes de estudio": "Sirven para validar pertenencia y limites.",
            "Que no se puede asumir": "No usar registros 2025 solamente; no usar malla como evidencia de cursado.",
            "Bloqueos actuales": "Fuente historica y deduplicacion acumulada no cerradas.",
            "Estado": "PENDIENTE",
        },
        {
            "Tema": "UNID_APROBADAS_TOTAL",
            "Detalle": "Unidades aprobadas acumuladas; aqui el instructivo permite incluir validacion de estudios/reconocimiento.",
            "Fuente requerida": "Historial completo con resultados aprobatorios y actos de validacion/homologacion.",
            "PROMEDIOS permite calcularlo": "PENDIENTE_DE_VALIDAR_COBERTURA_HISTORICA",
            "Planes de estudio": "Referencia de limite/tolerancia, no fuente de aprobacion.",
            "Que no se puede asumir": "No aplicar diccionario anual sin separar campo acumulado.",
            "Bloqueos actuales": "Regla de E/I para acumulado debe separarse del anual 19.",
            "Estado": "PENDIENTE",
        },
    ]
    return pd.DataFrame(rows)


def build_pipeline() -> pd.DataFrame:
    stages = [
        (
            "Input raw",
            "Fuentes F01-F07",
            "Inventarios y hashes",
            "gobernanza_integral_5809.py",
            "Rutas, hojas, columnas",
            "No modificar fuentes",
            "Hash SHA256 pre/post",
            "Ninguno si existen fuentes",
            "Reejecutar inventario",
        ),
        (
            "Limpieza",
            "CSV/TXT/XLSX raw",
            "Lectura normalizada",
            "gobernanza_integral_5809.py",
            "Encabezados y valores observados",
            "Normalizar nombres solo para inventario",
            "Conservar nombre original",
            "No corregir datos fuente",
            "Releer fuentes",
        ),
        (
            "Normalizacion",
            "5809, PROMEDIOS, MAPEO",
            "Llaves candidatas",
            "No autorizado para calculo final",
            "RUT, DV, CODCLI",
            "Normalizar RUT para cruce tecnico",
            "Cardinalidad por RUT",
            "Multiples CODCLI",
            "Volver a 06_LLAVES_GOBERNADAS",
        ),
        (
            "Mapeo de identidad",
            "MAPEO_IDENTIDAD_CODCLI",
            "CODCLI_LISTA gobernado",
            "No autorizado para calculo final",
            "CODCLI_LISTA, N_CODCLI",
            "Usar CODCLI_LISTA; bloquear N_CODCLI",
            "N_CODCLI no aparece como llave",
            "7 sin CODCLI_LISTA",
            "Resolver hoja 16_BLOQUEOS_ACTUALES",
        ),
        (
            "Filtro año 2025",
            "PROMEDIOS Hoja1",
            "Detalle 2025",
            "No autorizado para calculo final",
            "ANO",
            "ANO=2025",
            "No mezclar 2026",
            "122 sin registros 2025",
            "Auditar CODCLI/programa",
        ),
        (
            "Filtro CODCLI por programa",
            "CODCLI_LISTA + 5809 + 5810",
            "CODCLI de programa",
            "No autorizado para calculo final",
            "CODCLI, CODIGO_UNICO, PLAN_ESTUDIOS, CODCARR",
            "No sumar todos los CODCLI de un RUT",
            "Cardinalidad esperada",
            "Riesgo multi carrera",
            "Validacion funcional",
        ),
        (
            "Calculo semestres",
            "Detalle 2025",
            "16-17 solo si cerrado",
            "Bloqueado para definitivo",
            "PERIODO, CODRAMO",
            "PERIODO=1/2; otros en revision",
            "SI/NO y coherencia con cursadas",
            "71 diferencias 16-18/multiple",
            "Cerrar periodo y programa",
        ),
        (
            "Calculo cursadas",
            "Detalle 2025",
            "18 solo si cerrado",
            "Bloqueado para definitivo",
            "CODRAMO, ESTADO",
            "CODRAMO unico; NULL revision",
            "Cursadas >= aprobadas anual",
            "Pendientes paso 50",
            "Validacion funcional",
        ),
        (
            "Calculo aprobadas",
            "Detalle 2025 + diccionario",
            "19 solo si cerrado",
            "Bloqueado para definitivo",
            "ESTADO, DESCRIPCION_ESTADO",
            "A aprobado; E/I condicionados; R no; NULL revision",
            "Diccionario detectado",
            "223 diferencia solo aprobadas",
            "Decision funcional E/I/NULL",
        ),
        (
            "Auditoria",
            "Resultados candidatos",
            "Diferencias y pendientes",
            "Futuro script controlado",
            "Todas las columnas oficiales",
            "No declarar listo",
            "Manifest y hashes",
            "20-21 pendiente",
            "Revisar bloqueos",
        ),
        (
            "Salida no carga",
            "Gobernanza",
            "Excel, informe, manifest, script",
            "gobernanza_integral_5809.py",
            "Todas",
            "NO_LISTO_PARA_CARGA",
            "No SIES_READY",
            "Carga definitiva prohibida",
            "Solo tras cierre de gobernanza",
        ),
    ]
    return pd.DataFrame(
        [
            {
                "Etapa": s[0],
                "Entrada": s[1],
                "Salida": s[2],
                "Script": s[3],
                "Columnas usadas": s[4],
                "Reglas": s[5],
                "Validaciones": s[6],
                "Bloqueos": s[7],
                "Reanudacion": s[8],
            }
            for s in stages
        ]
    )


def build_transformations() -> Tuple[pd.DataFrame, pd.DataFrame]:
    autorizadas = [
        ("T01", "Normalizar RUT", "NUM_DOCUMENTO+DV / RUT+DIG", "RUT_NORMALIZADO", "Llave tecnica, no modifica fuente.", "Multiples CODCLI por RUT.", "Comparar DOC_PRECARGA_FULL vs RUT_INSTITUCIONAL.", "AUTORIZADA_GOBERNANZA"),
        ("T02", "Separar CODCLI_LISTA", "CODCLI_LISTA", "CODCLI candidato", "MAPEO_IDENTIDAD_CODCLI.", "Lista multiple requiere filtro por programa.", "No usar N_CODCLI.", "AUTORIZADA_CON_RESTRICCION"),
        ("T03", "Filtrar ANO=2025", "PROMEDIOS Hoja1", "Detalle 2025", "Año de referencia datos 2025.", "Mezclar año proceso 2026.", "Conteo registros por ANO.", "AUTORIZADA_GOBERNANZA"),
        ("T04", "Interpretar PERIODO=1/2", "PERIODO", "Semestre 1/2", "Instructivo pregunta primer/segundo semestre.", "PERIODO distinto de 1/2.", "Listar periodos observados.", "AUTORIZADA_CON_REVISION"),
        ("T05", "Contar CODRAMO unico para cursadas", "CODCLI filtrado + CODRAMO + ANO=2025", "UNIDADES_CURSADAS candidato", "Instructivo unidades cursadas.", "Duplicados por seccion/equivalencia.", "Deduplicar y auditar.", "AUTORIZADA_SOLO_TRAS_BLOQUEOS"),
        ("T06", "Contar CODRAMO unico con estado aprobatorio", "CODRAMO + ESTADO/DESCRIPCION_ESTADO", "UNIDADES_APROBADAS candidato", "Diccionario PROMEDIOS.", "E/I/NULL sin decision funcional.", "Aplicar hoja 05 y listar casos.", "AUTORIZADA_SOLO_TRAS_BLOQUEOS"),
        ("T07", "No usar N_CODCLI como CODCLI academico", "N_CODCLI", "BLOQUEO", "Incidente metodologico documentado.", "Confundir conteo con identificador.", "Buscar N_CODCLI en formulas.", "OBLIGATORIA"),
        ("T08", "No inferir estados A/E/I/R si existe diccionario", "ESTADO", "Diccionario aplicado", "PROMEDIOS contiene DESCRIPCION_ESTADO.", "Repetir error metodologico.", "Hoja 04/05.", "OBLIGATORIA"),
        ("T09", "Separar anual 16-19 de acumulado 20-21", "Campos 16-21", "Gobernanza separada", "Instructivo separa año 2025 y acumulado.", "Usar anual como acumulado.", "Hoja 09.", "OBLIGATORIA"),
        ("T10", "Registrar hash SHA256", "Fuentes y salidas", "Manifest", "Trazabilidad.", "No detectar cambios de fuente.", "Comparacion pre/post.", "AUTORIZADA"),
    ]
    prohibidas = [
        ("P01", "Usar N_CODCLI como CODCLI academico", "N_CODCLI es conteo, no llave.", "BLOQUEADA"),
        ("P02", "Inferir A/E/I/R sin revisar diccionario", "Existe ESTADO/DESCRIPCION_ESTADO en PROMEDIOS.", "BLOQUEADA"),
        ("P03", "Sumar todos los CODCLI de un RUT sin filtrar por programa", "Riesgo multi carrera.", "BLOQUEADA"),
        ("P04", "Usar estado academico para forzar unidades cero", "Estado academico no equivale a avance academico.", "BLOQUEADA"),
        ("P05", "Mezclar 2026 en calculo anual 2025", "Año proceso distinto de año de referencia.", "BLOQUEADA"),
        ("P06", "Presentar inferencia como regla oficial", "Separar dato observado, implementacion y decision interna.", "BLOQUEADA"),
        ("P07", "Presentar datos observados como norma", "PROMEDIOS no modifica instructivo.", "BLOQUEADA"),
        ("P08", "Modificar fuente original", "Regla de control obligatoria.", "BLOQUEADA"),
        ("P09", "Declarar SIES_READY sin cerrar pendientes", "Declaracion actual NO_LISTO_PARA_CARGA.", "BLOQUEADA"),
        ("P10", "Usar paso 48 como base de decision", "Paso 50 reemplaza inferencia previa sobre estados.", "BLOQUEADA"),
    ]
    df_aut = pd.DataFrame(
        [
            {
                "ID_TRANSFORMACION": row[0],
                "Descripcion": row[1],
                "Entrada": row[2],
                "Salida": row[3],
                "Regla/fuente": row[4],
                "Riesgo": row[5],
                "Prueba": row[6],
                "Estado": row[7],
            }
            for row in autorizadas
        ]
    )
    df_proh = pd.DataFrame(
        [
            {
                "ID_PROHIBICION": row[0],
                "Transformacion prohibida": row[1],
                "Justificacion": row[2],
                "Estado": row[3],
            }
            for row in prohibidas
        ]
    )
    return df_aut, df_proh


def build_validations() -> Tuple[pd.DataFrame, pd.DataFrame]:
    pre = [
        ("Fuente existe", "Cada fuente principal debe existir.", "OK", "Verificado por script."),
        ("Hash registrado", "Cada fuente debe tener SHA256.", "OK", "Hoja 01."),
        ("Hoja correcta identificada", "CSV/TXT/hojas Excel inventariadas.", "OK", "Hoja 02."),
        ("Columnas criticas detectadas", "CODCLI, ANO, PERIODO, CODRAMO, ESTADO, DESCRIPCION_ESTADO, CODCLI_LISTA.", "OK_PARCIAL", "Detectadas; uso definitivo bloqueado."),
        ("Diccionario detectado", "ESTADO/DESCRIPCION_ESTADO.", "OK", "Hojas 04 y 05."),
        ("Llave gobernada", "Llaves oficiales y tecnicas documentadas.", "OK_CON_RESTRICCIONES", "Hoja 06."),
        ("Cardinalidad validada", "RUT/CODCLI/programa.", "PENDIENTE", "Requiere validacion antes de calculo."),
        ("Año 2025 filtrado", "ANO=2025.", "PENDIENTE_CALCULO", "Documentado, no ejecutado para definitivo."),
        ("Período validado", "PERIODO 1/2 y otros.", "PENDIENTE", "PERIODO distinto de 1/2 queda en revision."),
        ("CODCLI correcto validado contra PROMEDIOS", "CODCLI_LISTA debe cruzar con PROMEDIOS.", "PENDIENTE", "122 sin registros y 7 sin CODCLI_LISTA."),
        ("No uso de columnas homonimas erradas", "N_CODCLI bloqueado.", "OK", "Hoja 12."),
        ("Regla de aprobadas validada", "A/E/I/R/NULL.", "PENDIENTE_FUNCIONAL", "E/I/NULL requieren cierre."),
        ("Resultado reproducible", "Script y manifest.", "OK", "Este expediente."),
        ("Pendientes visibles", "Bloqueos listados.", "OK", "Hoja 16."),
    ]
    post = [
        ("Total filas 5809 conserva 2371", "Debe verificarse despues de un calculo futuro.", "NO_EJECUTADO_EN_ESTE_PASO", "Este paso no genera carga."),
        ("Columnas oficiales conservan orden", "22 columnas oficiales.", "NO_EJECUTADO_EN_ESTE_PASO", "Hoja 07 documenta orden."),
        ("No se agregan columnas al CSV final", "Solo en eventual salida de carga.", "NO_EJECUTADO_EN_ESTE_PASO", "No se genero CSV final."),
        ("No se modifica precarga original", "Fuente original intacta.", "OK", "Hashes pre/post."),
        ("16-19 recalculadas solo si reglas gobernadas", "Bloqueado hasta cierre.", "BLOQUEADO", "No recalcular en este paso."),
        ("20-21 separadas", "Acumulado pendiente.", "OK", "Hoja 09."),
        ("Pendientes excluidos del definitivo", "No hay definitivo.", "OK", "Declaracion NO_LISTO_PARA_CARGA."),
        ("Diferencias documentadas", "Paso 50 y bloqueos.", "OK", "Hojas 15/16."),
        ("Manifest generado", "JSON.", "OK", "manifest_gobernanza_integral_5809.json."),
        ("Hash generado", "Fuentes y salidas.", "OK", "Manifest."),
        ("Declaracion carga correcta", "NO_LISTO_PARA_CARGA.", "OK", "Dictamen global."),
    ]
    return (
        pd.DataFrame(
            [
                {
                    "Validacion": r[0],
                    "Criterio": r[1],
                    "Estado": r[2],
                    "Observacion": r[3],
                }
                for r in pre
            ]
        ),
        pd.DataFrame(
            [
                {
                    "Validacion": r[0],
                    "Criterio": r[1],
                    "Estado": r[2],
                    "Observacion": r[3],
                }
                for r in post
            ]
        ),
    )


def build_incidents() -> pd.DataFrame:
    fecha = "2026-07-08"
    rows = [
        (
            "INC01",
            "Uso incorrecto de N_CODCLI como si fuera CODCLI academico.",
            "Flujo anterior de conciliacion/calculo.",
            "Confusion entre conteo de CODCLI y llave academica.",
            "Riesgo de cruzar registros academicos inexistentes o equivocados.",
            "Bloquear N_CODCLI y documentar CODCLI_LISTA como columna correcta.",
            "Checklist pre-calculo y prohibicion explicita.",
            "CORREGIDO_EN_GOBERNANZA",
        ),
        (
            "INC02",
            "Inferencia innecesaria de A/E/I/R pese a existir diccionario en PROMEDIOS.",
            "Paso 48 / inferencia previa de estados.",
            "No se inspecciono diccionario ESTADO/DESCRIPCION_ESTADO antes de interpretar codigos.",
            "Clasificacion de aprobadas quedo metodologicamente debil.",
            "Usar diccionario observado: A APROBADO, E CONVALIDACION, I HOMOLOGADO, R REPROBADO, NULL revision.",
            "Detectar diccionarios antes de cualquier inferencia.",
            "CORREGIDO_COMO_LECCION",
        ),
        (
            "INC03",
            "Riesgo de sumar todos los CODCLI del RUT sin filtrar por programa.",
            "Cruce RUT-CODCLI.",
            "RUT puede tener multiples CODCLI/programas.",
            "Unidades de otra carrera pueden contaminar 16-21.",
            "Gobernar llave RUT + CODIGO_UNICO + PLAN_ESTUDIOS y CODCLI_LISTA filtrado.",
            "Validacion de cardinalidad por programa.",
            "VIGENTE_COMO_RIESGO",
        ),
        (
            "INC04",
            "Riesgo de confundir estado academico con avance academico.",
            "Lectura de campos ESTADO_ACADEMICO/ESTADOACADEMICO.",
            "Campos de situacion academica no son equivalentes a resultados por ramo.",
            "Podria forzar unidades cero o vigencia sin regla.",
            "Bloquear uso de estado academico para calcular unidades.",
            "Separar catalogos de auditoria de formulas de avance.",
            "VIGENTE_COMO_RIESGO",
        ),
        (
            "INC05",
            "Riesgo de mezclar año proceso 2026 con datos academicos 2025.",
            "Filtros temporales.",
            "Proceso SIES 2026 usa año referencia datos 2025.",
            "Conteos anuales incorrectos.",
            "Documentar ANO=2025 para 16-19 y cierre academico 2025 para 20-21.",
            "Checklist de año/periodo antes de calculo.",
            "VIGENTE_COMO_RIESGO",
        ),
    ]
    return pd.DataFrame(
        [
            {
                "ID_INCIDENTE": r[0],
                "Descripcion": r[1],
                "Fecha": fecha,
                "Paso donde ocurrio": r[2],
                "Causa raiz": r[3],
                "Impacto": r[4],
                "Correccion": r[5],
                "Prevencion": r[6],
                "Estado": r[7],
            }
            for r in rows
        ]
    )


def read_recalc_summary() -> Dict[str, Any]:
    path = PATHS["F07_RECALCULO_50_EVIDENCIA"]["ruta"]
    rows = read_xlsx_rows(path, "00_DICTAMEN")
    if not rows:
        return {}
    return rows[0]


def build_blockers() -> pd.DataFrame:
    summary = read_recalc_summary()
    def val(key: str, fallback: Any) -> Any:
        return summary.get(key, fallback)

    rows = [
        (
            "B01",
            "423 pendientes 16-19 con diccionario estado PROMEDIOS",
            val("PENDIENTES", 423),
            "Paso 50 / hoja 00_DICTAMEN y 06_PENDIENTES_16_19",
            "Resolver diferencias y criterios E/I/NULL/programa antes de definitivo.",
            "Equipo funcional academico + datos",
            "NO",
            "Alto: no puede declararse listo.",
        ),
        (
            "B02",
            "223 diferencia solo aprobadas 19",
            val("DIFERENCIA_SOLO_APROBADAS_19", 223),
            "Paso 50",
            "Cerrar tratamiento de aprobadas, convalidaciones, homologaciones y NULL/blanco.",
            "Equipo funcional academico",
            "NO",
            "Alto: afecta UNIDADES_APROBADAS.",
        ),
        (
            "B03",
            "71 diferencia 16-18 o múltiple",
            val("DIFERENCIA_16_18_O_MULTIPLE", 71),
            "Paso 50",
            "Validar periodo, ramos, duplicados o mapeo programa.",
            "Datos + funcional",
            "NO",
            "Medio/alto: afecta presencia y cursadas.",
        ),
        (
            "B04",
            "122 sin registros 2025 para CODCLI_LISTA",
            val("SIN_REGISTROS_2025_PARA_CODCLI_LISTA", 122),
            "Paso 50",
            "Revisar CODCLI_LISTA, programa 2025 y cobertura PROMEDIOS.",
            "Datos institucionales",
            "NO",
            "Alto: ausencia de evidencia anual.",
        ),
        (
            "B05",
            "7 bloqueo sin CODCLI_LISTA",
            val("BLOQUEO_SIN_CODCLI_LISTA", 7),
            "MAPEO_IDENTIDAD_CODCLI / Paso 50",
            "Completar o resolver identidad academica antes de calcular.",
            "Datos institucionales",
            "NO",
            "Alto: sin llave academica.",
        ),
        (
            "B06",
            "20-21 acumulado pendiente/no cerrado",
            "PENDIENTE",
            "Gobernanza columnas 20-21",
            "Definir fuente historica completa, regla acumulada y limite por plan.",
            "Equipo funcional academico + datos",
            "NO",
            "Alto: separacion anual/acumulado no cerrada.",
        ),
    ]
    return pd.DataFrame(
        [
            {
                "ID_BLOQUEO": r[0],
                "Bloqueo": r[1],
                "Casos": r[2],
                "Fuente": r[3],
                "Requisito para resolver": r[4],
                "Responsable sugerido": r[5],
                "Puede resolverse automaticamente": r[6],
                "Riesgo": r[7],
            }
            for r in rows
        ]
    )


def build_resumption_rules() -> pd.DataFrame:
    rows = [
        ("R01", "No usar paso 48 como base de decision.", "El paso 48 correspondio a inferencia descartada para estados.", "OBLIGATORIA"),
        ("R02", "Paso 50 reemplaza paso 48 para estados.", "La evidencia tecnica valida el diccionario ESTADO/DESCRIPCION_ESTADO.", "OBLIGATORIA"),
        ("R03", "Antes de nuevo calculo, usar gobernanza 51.", "Este expediente es el punto de reanudacion metodologica.", "OBLIGATORIA"),
        ("R04", "Proximo paso tecnico permitido solo despues de revisar hoja 16_BLOQUEOS_ACTUALES.", "Los bloqueos impiden calculo definitivo.", "OBLIGATORIA"),
        ("R05", "No generar carga definitiva.", "Declaracion actual NO_LISTO_PARA_CARGA.", "OBLIGATORIA"),
        ("R06", "No modificar fuentes originales.", "Cualquier transformacion debe escribirse en derivado auditado.", "OBLIGATORIA"),
        ("R07", "No recalcular 20-21 hasta cerrar acumulado.", "Anual 16-19 no equivale a acumulado 20-21.", "OBLIGATORIA"),
    ]
    return pd.DataFrame(
        [
            {
                "ID_REGLA": r[0],
                "Regla de reanudacion": r[1],
                "Justificacion": r[2],
                "Estado": r[3],
            }
            for r in rows
        ]
    )


def build_executive_summary() -> pd.DataFrame:
    rows = [
        ("Que se encontro", "Las fuentes principales existen y fueron inventariadas; PROMEDIOS contiene diccionario ESTADO/DESCRIPCION_ESTADO."),
        ("Que estaba mal", "Se habia inferido A/E/I/R sin gobernar el diccionario interno; tambien existia riesgo de usar N_CODCLI como identificador."),
        ("Que se corrigio", "Se documento A/E/I/R/NULL, se bloqueo N_CODCLI, se goberno CODCLI_LISTA y se separo anual 16-19 de acumulado 20-21."),
        ("Que falta", "Cerrar bloqueos del paso 50, tratamiento funcional de E/I/NULL y regla/fuente acumulada 20-21."),
        ("Por que no esta listo para carga", "Hay 423 pendientes 16-19 y acumulado 20-21 no cerrado; el dictamen es NO_LISTO_PARA_CARGA."),
        ("Decision funcional requerida", "Confirmar aplicacion de convalidaciones/homologaciones en anual 2025 vs acumulado y resolver casos sin CODCLI_LISTA/sin registros 2025."),
    ]
    return pd.DataFrame([{"Tema": r[0], "Resumen": r[1]} for r in rows])


def build_sources_used(source_df: pd.DataFrame, output_paths: Dict[str, Path]) -> pd.DataFrame:
    rows = []
    for _, row in source_df.iterrows():
        rows.append(
            {
                "Tipo": "FUENTE_INPUT",
                "ID": row["ID_FUENTE"],
                "Ruta": row["Ruta"],
                "Hash SHA256": row["Hash SHA256"],
                "Observacion": row["Observacion"],
            }
        )
    for key, path in output_paths.items():
        rows.append(
            {
                "Tipo": "SALIDA_GOBERNANZA",
                "ID": key,
                "Ruta": str(path),
                "Hash SHA256": sha256_file(path) if path.exists() and path.is_file() else "",
                "Observacion": "Producto generado por este script.",
            }
        )
    return pd.DataFrame(rows)


def make_manifest_readable(manifest: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    simple_keys = [
        "proceso",
        "subproyecto",
        "anio_proceso",
        "anio_referencia_datos",
        "timestamp",
        "estado_carga",
        "fuentes_originales_modificadas",
        "dictamen_global",
        "proximo_paso_permitido",
    ]
    for key in simple_keys:
        rows.append({"Seccion": "GLOBAL", "Clave": key, "Valor": manifest.get(key, "")})
    for key, value in manifest.get("conteos", {}).items():
        rows.append({"Seccion": "CONTEOS", "Clave": key, "Valor": value})
    for requested, physical in manifest.get("equivalencias_hojas_excel", {}).items():
        rows.append({"Seccion": "HOJAS_EXCEL", "Clave": requested, "Valor": physical})
    for key, value in manifest.get("archivos", {}).items():
        rows.append({"Seccion": "ARCHIVOS", "Clave": key, "Valor": value})
    return pd.DataFrame(rows)


def build_markdown(
    dictamen_df: pd.DataFrame,
    incidents_df: pd.DataFrame,
    blockers_df: pd.DataFrame,
    manifest: Dict[str, Any],
) -> str:
    lines = [
        "# Gobernanza integral inputs, columnas y transformaciones 5809",
        "",
        f"Proceso: **{PROCESO}**",
        f"Subproyecto: **{SUBPROYECTO}**",
        f"Año proceso: **{ANIO_PROCESO}**",
        f"Año referencia datos: **{ANIO_REFERENCIA_DATOS}**",
        f"Declaración carga: **{DECLARACION_CARGA}**",
        "",
        "## Dictamen global",
        "",
        markdown_table(dictamen_df),
        "",
        "## Hallazgos principales",
        "",
        "- Se inventariaron las fuentes, hojas y columnas principales antes de usar columnas para cálculo.",
        "- Se detectó y documentó el diccionario `ESTADO` / `DESCRIPCION_ESTADO` de PROMEDIOS.",
        "- `N_CODCLI` queda bloqueado como CODCLI académico; `CODCLI_LISTA` queda documentado como columna correcta de mapeo.",
        "- Las columnas anuales 16-19 se separan de las acumuladas 20-21.",
        "- Paso 48 queda descartado como inferencia para estados; paso 50 reemplaza esa base técnica.",
        "- El expediente no genera carga, no declara listo y no modifica fuentes originales.",
        "",
        "## Incidentes metodológicos",
        "",
        markdown_table(incidents_df),
        "",
        "## Bloqueos actuales",
        "",
        markdown_table(blockers_df),
        "",
        "## Equivalencias de hojas Excel",
        "",
        "Excel limita los nombres de hoja a 31 caracteres. Cuando un nombre solicitado excede ese límite, el contenido obligatorio se conserva en una hoja física abreviada y esta equivalencia queda en el manifest.",
        "",
        markdown_table(pd.DataFrame(
            [
                {"Nombre solicitado": req, "Nombre fisico Excel": phys}
                for req, phys in REQUESTED_SHEETS.items()
            ]
        )),
        "",
        "## Qué sigue pendiente",
        "",
        "- Revisar la hoja `16_BLOQUEOS_ACTUALES` antes de cualquier cálculo.",
        "- Confirmar tratamiento funcional de convalidación (`E`), homologación (`I`) y `NULL/blanco` por campo.",
        "- Cerrar regla y fuente de acumulados 20-21.",
        "- Solo después de cerrar gobernanza, autorizar un cálculo controlado no definitivo; la carga final sigue prohibida.",
        "",
        "## Archivos",
        "",
    ]
    for key, value in manifest.get("archivos", {}).items():
        lines.append(f"- {key}: `{value}`")
    lines.append("")
    return "\n".join(lines)


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        max_col = min(ws.max_column, 40)
        for idx in range(1, max_col + 1):
            letter = get_column_letter(idx)
            max_len = 12
            for cell in ws[letter][: min(ws.max_row, 200)]:
                text = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(text), 80))
            ws.column_dimensions[letter].width = min(max_len + 2, 60)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    wb.close()


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display_value(v, 80) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * widths[i] for i in range(len(cols))) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display_value(row[col], 80).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def write_excel(output_path: Path, dataframes: OrderedDict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for requested_name, df in dataframes.items():
            physical_name = REQUESTED_SHEETS[requested_name]
            df.to_excel(writer, sheet_name=physical_name, index=False)
    style_workbook(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--timestamp", default=datetime.now().strftime("%Y%m%d_%H%M%S"))
    args = parser.parse_args()

    output_dir = BASE_51 / f"GOBERNANZA_INTEGRAL_5809_{args.timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    desktop_dir = Path.home() / "Desktop" / f"AVANCE_CURRICULAR_2026_GOBERNANZA_INTEGRAL_5809_{args.timestamp}"
    desktop_dir.mkdir(parents=True, exist_ok=True)

    excel_path = output_dir / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
    md_path = output_dir / "INFORME_GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.md"
    manifest_path = output_dir / "manifest_gobernanza_integral_5809.json"
    script_path = output_dir / "gobernanza_integral_5809.py"
    current_script = Path(__file__).resolve()
    if current_script != script_path.resolve():
        shutil.copy2(current_script, script_path)

    source_hashes_before = {source_id: sha256_file(meta["ruta"]) for source_id, meta in PATHS.items()}

    source_df, sheets_df, columns_df, source_hashes_inventory, headers_5809 = build_inventories()
    dictionaries_df, estado_df = build_dictionaries()
    keys_df = build_keys()
    mapeo_5809_df = build_mapeo_5809(headers_5809)
    gov_16_19_df = build_governance_16_19()
    gov_20_21_df = build_governance_20_21()
    pipeline_df = build_pipeline()
    transformations_df, prohibited_df = build_transformations()
    validations_pre_df, validations_post_df = build_validations()
    incidents_df = build_incidents()
    blockers_df = build_blockers()
    resumption_df = build_resumption_rules()
    executive_df = build_executive_summary()

    total_gobernadas = int((columns_df["Estado gobernanza"] == "GOBERNADA").sum())
    total_no_gobernadas = int(len(columns_df) - total_gobernadas)
    total_llaves_gobernadas = int(keys_df["Estado gobernanza"].str.contains("GOBERNADA", na=False).sum())
    total_diccionarios = int(dictionaries_df["ID_DICCIONARIO"].nunique())
    total_transformaciones = int(len(transformations_df))
    total_bloqueos = int(len(blockers_df))

    dictamen_df = pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Estado de carga": DECLARACION_CARGA,
                "Dictamen global": (
                    "Gobernanza generada. No se autoriza carga ni recalculo definitivo hasta cerrar "
                    "bloqueos 16-21 y decisiones funcionales."
                ),
                "Fuentes originales modificadas": "NO",
                "Total fuentes inventariadas": len(source_df),
                "Total hojas inventariadas": len(sheets_df),
                "Total columnas inventariadas": len(columns_df),
                "Total columnas gobernadas": total_gobernadas,
                "Total columnas no gobernadas": total_no_gobernadas,
                "Total llaves gobernadas": total_llaves_gobernadas,
                "Total diccionarios detectados": total_diccionarios,
                "Total transformaciones permitidas": total_transformaciones,
                "Total bloqueos": total_bloqueos,
                "Proximo paso permitido": (
                    "Revisar 16_BLOQUEOS_ACTUALES, validar E/I/NULL y cerrar acumulados 20-21; "
                    "no generar carga definitiva."
                ),
            }
        ]
    )

    placeholder_outputs = {
        "Excel": excel_path,
        "Informe": md_path,
        "Manifest": manifest_path,
        "Script": script_path,
        "Carpeta escritorio": desktop_dir,
    }
    sources_used_df = build_sources_used(source_df, placeholder_outputs)

    preliminary_manifest = {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "timestamp": args.timestamp,
        "estado_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "dictamen_global": dictamen_df.loc[0, "Dictamen global"],
        "proximo_paso_permitido": dictamen_df.loc[0, "Proximo paso permitido"],
        "conteos": {
            "total_fuentes_inventariadas": len(source_df),
            "total_hojas_inventariadas": len(sheets_df),
            "total_columnas_inventariadas": len(columns_df),
            "total_columnas_gobernadas": total_gobernadas,
            "total_columnas_no_gobernadas": total_no_gobernadas,
            "total_llaves_gobernadas": total_llaves_gobernadas,
            "total_diccionarios_detectados": total_diccionarios,
            "total_transformaciones_permitidas": total_transformaciones,
            "total_bloqueos": total_bloqueos,
        },
        "equivalencias_hojas_excel": REQUESTED_SHEETS,
        "archivos": {k: str(v) for k, v in placeholder_outputs.items()},
    }
    manifest_readable_df = make_manifest_readable(preliminary_manifest)

    dataframes = OrderedDict(
        [
            ("00_DICTAMEN_GLOBAL", dictamen_df),
            ("01_FUENTES_INVENTARIO", source_df),
            ("02_HOJAS_INVENTARIO", sheets_df),
            ("03_COLUMNAS_INVENTARIO", columns_df),
            ("04_DICCIONARIOS_DETECTADOS", dictionaries_df),
            ("05_DICCIONARIO_ESTADO_PROMEDIOS", estado_df),
            ("06_LLAVES_GOBERNADAS", keys_df),
            ("07_MAPEO_5809_COLUMNAS_OFICIALES", mapeo_5809_df),
            ("08_GOBERNANZA_COLUMNAS_16_19", gov_16_19_df),
            ("09_GOBERNANZA_COLUMNAS_20_21", gov_20_21_df),
            ("10_FLUJO_RAW_A_PROCESADO", pipeline_df),
            ("11_TRANSFORMACIONES_AUTORIZADAS", transformations_df),
            ("12_TRANSFORMACIONES_PROHIBIDAS", prohibited_df),
            ("13_VALIDACIONES_OBLIGATORIAS_PRE_CALCULO", validations_pre_df),
            ("14_VALIDACIONES_OBLIGATORIAS_POST_CALCULO", validations_post_df),
            ("15_INCIDENTES_METODOLOGICOS", incidents_df),
            ("16_BLOQUEOS_ACTUALES", blockers_df),
            ("17_REGLAS_DE_REANUDACION", resumption_df),
            ("18_RESUMEN_EJECUTIVO", executive_df),
            ("19_FUENTES", sources_used_df),
            ("20_MANIFEST_LEGIBLE", manifest_readable_df),
        ]
    )

    write_excel(excel_path, dataframes)

    manifest = dict(preliminary_manifest)
    manifest["fuentes"] = source_df.to_dict(orient="records")
    manifest["bloqueos"] = blockers_df.to_dict(orient="records")
    manifest["incidentes_metodologicos"] = incidents_df.to_dict(orient="records")
    manifest["reglas_de_reanudacion"] = resumption_df.to_dict(orient="records")
    manifest["hashes_fuentes_antes"] = source_hashes_before
    manifest["hashes_fuentes_inventario"] = source_hashes_inventory

    md_text = build_markdown(dictamen_df, incidents_df, blockers_df, manifest)
    md_path.write_text(md_text, encoding="utf-8")

    source_hashes_after = {source_id: sha256_file(meta["ruta"]) for source_id, meta in PATHS.items()}
    fuentes_modificadas = "SI" if source_hashes_before != source_hashes_after else "NO"
    dictamen_df.loc[0, "Fuentes originales modificadas"] = fuentes_modificadas
    manifest["fuentes_originales_modificadas"] = fuentes_modificadas
    manifest["hashes_fuentes_despues"] = source_hashes_after
    manifest["archivos"] = {
        "Excel": str(excel_path),
        "Informe": str(md_path),
        "Manifest": str(manifest_path),
        "Script": str(script_path),
        "Carpeta escritorio": str(desktop_dir),
    }
    manifest["hashes_salidas"] = {
        "Excel": sha256_file(excel_path),
        "Informe": sha256_file(md_path),
        "Script": sha256_file(script_path),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    for path in [excel_path, md_path, manifest_path, script_path]:
        shutil.copy2(path, desktop_dir / path.name)

    print("GOBERNANZA INTEGRAL 5809 — GENERADA")
    print(f"Fuentes originales modificadas: {fuentes_modificadas}")
    print(f"Declaración carga: {DECLARACION_CARGA}")
    print()
    print("DICTAMEN GLOBAL")
    print(markdown_table(dictamen_df))
    print()
    print("INCIDENTES METODOLOGICOS")
    print(markdown_table(incidents_df[["ID_INCIDENTE", "Descripcion", "Estado"]]))
    print()
    print("BLOQUEOS ACTUALES")
    print(markdown_table(blockers_df[["ID_BLOQUEO", "Bloqueo", "Casos", "Riesgo"]]))
    print()
    print("ARCHIVOS")
    print(f"Excel: {excel_path}")
    print(f"Informe: {md_path}")
    print(f"Manifest: {manifest_path}")
    print(f"Script: {script_path}")
    print(f"Carpeta escritorio: {desktop_dir}")


if __name__ == "__main__":
    main()
