#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Hito 67: gobernanza especifica de PERIODO raw para Avance Curricular
SIES 2026, archivo 5809, columnas 16-18.

Este hito es documental, exploratorio, trazable y reproducible. No modifica
fuentes originales, no corrige datos, no genera carga, no genera SIES_READY,
no recalcula 20-21 y no aplica reglas de PERIODO raw.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Gobernanza PERIODO raw para columnas 16, 17 y 18 del archivo 5809 Matricula Avance Curricular"
ANIO_PROCESO = 2026
ANIO_REFERENCIA = 2025
ESTADO = "EXPLORATORIO / GOBERNANZA"
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_OUT = REPO / "avance_curricular_2026" / "67_gobernanza_periodo_raw_5809"

INSTRUCTIVO = REPO / "avance_curricular_2026/00_fuentes_congeladas/CARGA_CONGELADA_20260626_005826/originales/Instructivo_Avance Curricular SIES - 2026.txt"
PROMEDIOS = REPO / "avance_curricular_2026/25_actualizacion_fuente_promedios/CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/00_FUENTE_CONGELADA/PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
H55_DIR = REPO / "avance_curricular_2026/55_auditoria_diferencia_16_18_multiple_5809/AUDITORIA_DIF_16_18_MULTIPLE_20260708_172433"
H55_XLSX = H55_DIR / "AUDITORIA_DIFERENCIA_16_18_MULTIPLE_5809.xlsx"
H55_SCRIPT = H55_DIR / "auditoria_diferencia_16_18_multiple_5809.py"
H55_MD = H55_DIR / "INFORME_AUDITORIA_DIFERENCIA_16_18_MULTIPLE_5809.md"
H63_DIR = REPO / "avance_curricular_2026/63_terminal_revision_38_decision_funcional_restante_5809/REVISION_38_DECISION_FUNCIONAL_RESTANTE_20260708_221125"
H63_XLSX = H63_DIR / "REVISION_38_DECISION_FUNCIONAL_RESTANTE_5809.xlsx"
H64_DIR = REPO / "avance_curricular_2026/64_exploracion_gobernanza_transversal_mu_avance/EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE_20260708_221914"
H64_XLSX = H64_DIR / "EXPLORACION_GOBERNANZA_TRANSVERSAL_MU_AVANCE.xlsx"
H66_DIR = REPO / "avance_curricular_2026/66_correccion_dictamen_post_h64_16_19_5809/CORRECCION_DICTAMEN_POST_H64_16_19_20260708_230001"
H66_XLSX = H66_DIR / "CORRECCION_DICTAMEN_POST_H64_16_19_5809.xlsx"
H66_PROMPT = H66_DIR / "PROMPT_GOBERNANZA_PERIODO_RAW_5809.md"
CONTROL_FOR_ING = REPO / "control/for_ing_act_trace_long.tsv"
CONTROL_CODCARPR = REPO / "control/gob_codcarpr_anioingreso_long.tsv"
CONTROL_CAMPOS_ING = REPO / "control/campos_ing_trace_long.tsv"
MU_CODE = REPO / "codigo_gobernanza_v2.py"
MU_AUDIT = REPO / "resultados/run_and_validate_oficial_2026-04-09_sin_puente/auditoria_maestra.md"

OUTPUT_FILES = {
    "excel": "GOBERNANZA_PERIODO_RAW_5809.xlsx",
    "informe": "INFORME_GOBERNANZA_PERIODO_RAW_5809.md",
    "manifest": "manifest_gobernanza_periodo_raw_5809.json",
    "script": "gobernanza_periodo_raw_5809.py",
    "borrador": "BORRADOR_DECISION_FUNCIONAL_PERIODO_RAW_5809.md",
}

SHEETS = OrderedDict(
    [
        ("00_DICTAMEN_GLOBAL", "00_DICTAMEN_GLOBAL"),
        ("01_FUENTES_REVISADAS", "01_FUENTES_REVISADAS"),
        ("02_EXTRACTOS_INSTRUCTIVO", "02_EXTRACTOS_INSTRUCTIVO"),
        ("03_CAMPOS_PERIODO_DETECTADOS", "03_CAMPOS_PERIODO_DETECT"),
        ("04_PROMEDIOS_PERIODO_RAW", "04_PROMEDIOS_PERIODO_RAW"),
        ("05_17_CASOS_PERIODO_NO_1_2", "05_17_CASOS_PERIODO_NO_1_2"),
        ("06_DETALLE_RAMO_PERIODO_17", "06_DETALLE_RAMO_PERIODO_17"),
        ("07_ESCENARIOS_DE_MAPEO", "07_ESCENARIOS_DE_MAPEO"),
        ("08_MATRIZ_DECISION_PERIODO", "08_MATRIZ_DECISION_PERIODO"),
        ("09_MAPEOS_NO_APLICAR", "09_MAPEOS_NO_APLICAR"),
        ("10_RECOMENDACION_FUNCIONAL", "10_RECOMENDACION_FUNCIONAL"),
        ("11_BORRADOR_DECISION_FUNCIONAL", "11_BORRADOR_DECISION_FUNC"),
        ("12_BLOQUEOS_Y_PENDIENTES", "12_BLOQUEOS_Y_PENDIENTES"),
        ("13_FUENTES", "13_FUENTES"),
        ("14_MANIFEST_LEGIBLE", "14_MANIFEST_LEGIBLE"),
    ]
)

TERMS_PERIODO = ["PERIODO", "PERIODOMATRICULA", "PERIODOINGRESO"]
TERMS_SEMESTRE = ["SEMESTRE", "SEM_INGRESO", "SEM_ING", "CURSO_1ER_SEM", "CURSO_2DO_SEM"]
FIELD_TERMS = TERMS_PERIODO + TERMS_SEMESTRE
EXCLUDE_DIRS = {".git", ".venv", "__pycache__", ".mypy_cache", ".pytest_cache"}

DICTAMEN_GLOBAL = (
    "No se detecta regla oficial que defina el significado de PERIODO raw 1,2,3,4,5 en PROMEDIOS "
    "para Avance Curricular 5809. El instructivo gobierna los conceptos primer semestre, segundo "
    "semestre y unidades cursadas 2025, pero no homologa esos conceptos al campo raw PERIODO. "
    "Los periodos 1 y 2 tienen uso tecnico trazable como candidatos, no como regla final; "
    "PERIODO 3, 4 y 5 permanecen bloqueados. Los 17 casos PERIODO_NO_1_2 siguen bloqueados "
    "hasta decision funcional/documental especifica."
)

BLOQUEO_H64 = (
    "Hito 64: PERIODO raw 3/4/5 no confirmado como semestre para Avance ni MU; "
    "PERIODO 2/3 como segundo semestre queda pendiente/bloqueado; no aplicar transversalmente sin confirmacion."
)

BLOQUEO_H66 = (
    "Hito 66: los 17 casos PERIODO_NO_1_2 permanecen bloqueados; no se aplica 1=primer semestre "
    "ni 2/3=segundo semestre sin respaldo documental o validacion funcional formal."
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    if text.endswith(".0"):
        return text[:-2]
    return text


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def display(value: Any, max_len: int = 110) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def md_cell(value: Any, max_len: int = 110) -> str:
    return display(value, max_len).replace("|", "\\|").replace("\n", "<br>")


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [md_cell(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append("| " + " | ".join(md_cell(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols)) + " |")
    return "\n".join([header, sep] + body)


def read_xlsx(path: Path, sheet: str, **kwargs: Any) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str, **kwargs).fillna("")


def read_tsv(path: Path, nrows: int | None = None) -> pd.DataFrame:
    return pd.read_csv(path, sep="\t", dtype=str, nrows=nrows, encoding="utf-8", on_bad_lines="skip").fillna("")


def is_under(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def safe_rel(path: Path) -> str:
    try:
        return str(path.relative_to(REPO))
    except ValueError:
        return str(path)


def classify_source(path: Path) -> Tuple[str, str]:
    text = str(path).lower()
    suffix = path.suffix.lower()
    if path == INSTRUCTIVO:
        return "instructivo oficial", "A"
    if path == PROMEDIOS:
        return "fuente academica observada", "B"
    if suffix == ".py":
        return "implementacion tecnica", "C"
    if "hito 64" in text or "64_exploracion" in text or "hito 66" in text or "66_correccion" in text:
        return "decision/exploracion interna", "D/E"
    if "55_auditoria" in text or "63_terminal" in text:
        return "auditoria tecnica", "C/D"
    if suffix in {".xlsx", ".tsv", ".csv"}:
        return "dato observado/control", "B/C"
    if suffix in {".md", ".json"}:
        return "reporte/manifest/evidencia", "C/D/E"
    return "evidencia", "E"


def scan_text_flags(path: Path, max_bytes: int = 2_000_000) -> Tuple[str, str, str, str]:
    if not path.exists() or not path.is_file():
        return "NO", "NO", "NO", ""
    if path.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
        return "NO_LEIDO_TEXTO", "NO_LEIDO_TEXTO", "NO_LEIDO_TEXTO", "Excel revisado por hojas/columnas."
    try:
        with path.open("rb") as fh:
            raw = fh.read(max_bytes)
        text = raw.decode("utf-8", errors="ignore")
    except Exception as exc:
        return "NO", "NO", "NO", f"No se pudo leer como texto: {exc}"
    upper = text.upper()
    contains_periodo = any(term in upper for term in TERMS_PERIODO)
    contains_semestre = any(term in upper for term in TERMS_SEMESTRE)
    mapping_patterns = [
        r"PERIODO\s*=\s*3.*SEGUNDO",
        r"2\s*,\s*3.*\]\)\s*=\s*2",
        r"2/3\s*=\s*SEGUNDO",
        r"PERIODO.*2\s*/\s*3",
        r"3\s*->\s*2",
        r"3\s*→\s*2",
        r"EQUIVALENTE OPERACIONAL DE SEGUNDO SEMESTRE",
    ]
    contains_mapping = any(re.search(pattern, upper, flags=re.S) for pattern in mapping_patterns)
    obs = "Busqueda textual acotada a 2MB; suficiente para evidencia de encabezados/codigo cercano." if len(raw) >= max_bytes else ""
    return ("SI" if contains_periodo else "NO", "SI" if contains_semestre else "NO", "SI" if contains_mapping else "NO", obs)


def discover_local_files(output_dir: Path) -> List[Path]:
    files: List[Path] = []
    patterns = ["*.tsv", "control/*.tsv", "**/*.py", "**/*.md", "**/*.json"]
    for pattern in patterns:
        for path in REPO.glob(pattern):
            if not path.is_file():
                continue
            parts = set(path.parts)
            if parts & EXCLUDE_DIRS:
                continue
            if is_under(path, output_dir.parent):
                continue
            files.append(path)
    return sorted(set(files), key=lambda p: str(p).lower())


def build_sources(output_dir: Path) -> pd.DataFrame:
    mandatory = [
        INSTRUCTIVO,
        PROMEDIOS,
        H55_DIR,
        H55_XLSX,
        H55_SCRIPT,
        H55_MD,
        H63_DIR,
        H63_XLSX,
        H64_DIR,
        H64_XLSX,
        H66_DIR,
        H66_XLSX,
        H66_PROMPT,
        CONTROL_FOR_ING,
        CONTROL_CODCARPR,
        CONTROL_CAMPOS_ING,
        MU_CODE,
        MU_AUDIT,
    ]
    local = discover_local_files(output_dir)
    all_paths = sorted(set(mandatory + local), key=lambda p: str(p).lower())
    rows = []
    for path in all_paths:
        exists = path.exists()
        source_type, level = classify_source(path)
        contains_periodo, contains_semestre, contains_mapping, scan_obs = scan_text_flags(path) if exists and path.is_file() else ("NO", "NO", "NO", "")
        rows.append(
            {
                "Ruta": str(path),
                "Existe SI/NO": "SI" if exists else "NO",
                "Tipo": path.suffix.lower().lstrip(".") if path.is_file() else "carpeta",
                "Clasificación fuente": source_type,
                "Nivel respaldo A/B/C/D/E": level,
                "Contiene PERIODO SI/NO": contains_periodo,
                "Contiene SEMESTRE SI/NO": contains_semestre,
                "Contiene mapeo explícito SI/NO": contains_mapping,
                "Hash SHA256": sha256_file(path) if exists and path.is_file() else "",
                "Observación": scan_obs,
            }
        )
    return pd.DataFrame(rows)


def build_extractos_instructivo() -> pd.DataFrame:
    lines = INSTRUCTIVO.read_text(encoding="utf-8", errors="replace").splitlines()
    specs = [
        (45, "Alcance carga matricula avance", "Columnas 16-18", "A", "Debe informar actividad en primer/segundo semestre y unidades cursadas/aprobadas 2025."),
        (47, "Primer y segundo semestre", "CURSO_1ER_SEM / CURSO_2DO_SEM", "A", "Concepto oficial; no define PERIODO raw."),
        (50, "Ano academico 2025", "UNIDADES_CURSADAS", "A", "Concepto anual oficial."),
        (155, "Presencia estudiante carrera", "CURSO_1ER_SEM / CURSO_2DO_SEM", "A", "Presencia en carrera informada."),
        (157, "Primer y segundo semestre", "CURSO_1ER_SEM / CURSO_2DO_SEM", "A", "Gobierna concepto de semestre; no codifica PERIODO raw."),
        (158, "Unidades Cursadas", "UNIDADES_CURSADAS", "A", "Inicio definicion unidades cursadas."),
        (160, "Unidades Cursadas", "UNIDADES_CURSADAS", "A", "Cursada con resultado aprobacion o reprobacion."),
        (162, "Ultimo ano academico 2025", "UNIDADES_CURSADAS", "A", "Distingue anual 2025."),
        (231, "Campos faltantes precarga", "16-18", "A", "Campos a llenar: presencia y avance 2025."),
        (232, "Presencia 1er y 2do semestre", "CURSO_1ER_SEM / CURSO_2DO_SEM", "A", "Repite alcance oficial."),
        (427, "SEM_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL", "A", "Semestre oficial de ingreso, no PERIODO raw PROMEDIOS."),
        (429, "SEM_INGRESO_CARRERA_ACTUAL codigos", "SEM_INGRESO_CARRERA_ACTUAL", "A", "Valores 1/2 para campo de ingreso."),
        (436, "SEM_INGRESO_CARRERA_ORIGEN", "SEM_INGRESO_CARRERA_ORIGEN", "A", "Semestre oficial de origen, no PERIODO raw PROMEDIOS."),
        (438, "SEM_INGRESO_CARRERA_ORIGEN codigos", "SEM_INGRESO_CARRERA_ORIGEN", "A", "Valores 1/2 para campo de ingreso."),
        (440, "CURSO_1ER_SEM", "CURSO_1ER_SEM", "A", "Definicion campo 16."),
        (441, "CURSO_1ER_SEM", "CURSO_1ER_SEM", "A", "Primer semestre del ano 2025."),
        (449, "CURSO_2DO_SEM", "CURSO_2DO_SEM", "A", "Definicion campo 17."),
        (450, "CURSO_2DO_SEM", "CURSO_2DO_SEM", "A", "Segundo semestre del ano 2025."),
        (454, "UNIDADES_CURSADAS", "UNIDADES_CURSADAS", "A", "Definicion campo 18."),
        (457, "UNIDADES_CURSADAS", "UNIDADES_CURSADAS", "A", "Efectivamente durante ano academico 2025."),
        (642, "Validacion CURSO_1ER_SEM", "CURSO_1ER_SEM", "A", "Valores permitidos SI/NO."),
        (647, "Validacion CURSO_1ER_SEM", "CURSO_1ER_SEM / UNIDADES_CURSADAS", "A", "Si SI, unidades cursadas > 0."),
        (648, "Validacion CURSO_2DO_SEM", "CURSO_2DO_SEM / UNIDADES_CURSADAS", "A", "Si SI, unidades cursadas > 0."),
        (665, "Validacion CURSO_2DO_SEM", "CURSO_2DO_SEM", "A", "Valores permitidos SI/NO."),
        (666, "Validacion ambos NO", "CURSO_1ER_SEM / CURSO_2DO_SEM / UNIDADES_CURSADAS", "A", "Si ambos NO, UNIDADES_CURSADAS debe ser 0."),
    ]
    rows = []
    for line_no, tema, campo, level, obs in specs:
        rows.append(
            {
                "Texto exacto": lines[line_no - 1].strip() if 1 <= line_no <= len(lines) else "",
                "Campo afectado": campo,
                "Nivel respaldo A": level,
                "Linea/seccion": f"Linea {line_no}",
                "Tema": tema,
                "Observación": obs,
            }
        )
    return pd.DataFrame(rows)


def value_counts_text(series: pd.Series, limit: int = 20) -> str:
    vc = series.map(clean).value_counts(dropna=False).head(limit)
    return "; ".join(f"{idx}={val}" for idx, val in vc.items())


def build_campos_periodo_detectados(sources: pd.DataFrame) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    # Excel principales.
    excel_paths = [PROMEDIOS, H55_XLSX, H63_XLSX, H64_XLSX, H66_XLSX]
    for path in excel_paths:
        if not path.exists():
            continue
        xl = pd.ExcelFile(path)
        for sheet in xl.sheet_names:
            try:
                header = pd.read_excel(path, sheet_name=sheet, nrows=0)
                columns = [str(c) for c in header.columns]
            except Exception:
                continue
            for col in columns:
                upper = col.upper()
                if not any(term in upper for term in FIELD_TERMS):
                    continue
                values = ""
                counts = ""
                try:
                    df_col = pd.read_excel(path, sheet_name=sheet, usecols=[col], dtype=str).fillna("")
                    values = " | ".join(sorted(set(clean(v) for v in df_col[col].tolist() if clean(v)))[:30])
                    counts = value_counts_text(df_col[col])
                except Exception as exc:
                    values = "NO_LEIDO"
                    counts = f"Error lectura valores: {exc}"
                rows.append(
                    {
                        "Archivo": str(path),
                        "Hoja si aplica": sheet,
                        "Columna": col,
                        "Valores únicos": values,
                        "Conteo por valor": counts,
                        "Uso probable": infer_usage(col, path),
                        "Nivel respaldo": classify_source(path)[1],
                        "Observación": "Campo detectado en Excel.",
                    }
                )
    # TSV top-level y control.
    for path in sorted(set(list(REPO.glob("*.tsv")) + list((REPO / "control").glob("*.tsv"))), key=lambda p: str(p).lower()):
        if not path.exists():
            continue
        try:
            header = path.open("r", encoding="utf-8", errors="ignore").readline().strip().split("\t")
        except Exception:
            header = []
        for col in header:
            if not any(term in col.upper() for term in FIELD_TERMS):
                continue
            values = ""
            counts = ""
            try:
                df = read_tsv(path, nrows=5000)
                if col in df.columns:
                    values = " | ".join(sorted(set(clean(v) for v in df[col].tolist() if clean(v)))[:30])
                    counts = value_counts_text(df[col])
            except Exception as exc:
                values = "NO_LEIDO"
                counts = f"Error lectura TSV: {exc}"
            rows.append(
                {
                    "Archivo": str(path),
                    "Hoja si aplica": "TSV",
                    "Columna": col,
                    "Valores únicos": values,
                    "Conteo por valor": counts,
                    "Uso probable": infer_usage(col, path),
                    "Nivel respaldo": classify_source(path)[1],
                    "Observación": "TSV local/control; no reemplaza instructivo.",
                }
            )
    # Evidencia textual/codigo con mapeo explicito.
    for path in [MU_CODE, MU_AUDIT, H55_SCRIPT, H55_MD, H66_PROMPT]:
        if not path.exists():
            continue
        cp, cs, cm, obs = scan_text_flags(path)
        if cp == "SI" or cs == "SI":
            rows.append(
                {
                    "Archivo": str(path),
                    "Hoja si aplica": "texto/codigo",
                    "Columna": "(no aplica)",
                    "Valores únicos": "",
                    "Conteo por valor": "",
                    "Uso probable": "Implementacion o reporte; no regla oficial.",
                    "Nivel respaldo": classify_source(path)[1],
                    "Observación": f"Contiene PERIODO={cp}; SEMESTRE={cs}; mapeo explicito={cm}. {obs}",
                }
            )
    return pd.DataFrame(rows)


def infer_usage(col: str, path: Path) -> str:
    upper = col.upper()
    ptext = str(path).lower()
    if upper == "PERIODO" and path == PROMEDIOS:
        return "PERIODO raw academico PROMEDIOS; objeto central del hito."
    if "PERIODOMATRICULA" in upper:
        return "Periodo de matricula; no equivale automaticamente a semestre Avance 5809."
    if "PERIODOINGRESO" in upper:
        return "Periodo de ingreso; no equivale automaticamente a PERIODO raw de ramos."
    if "SEM_INGRESO" in upper or "SEM_ING" in upper:
        return "Semestre oficial/derivado de ingreso; concepto distinto de PERIODO raw de actividad."
    if "CURSO_1ER_SEM" in upper:
        return "Columna 16 oficial/resultado auditado."
    if "CURSO_2DO_SEM" in upper:
        return "Columna 17 oficial/resultado auditado."
    if "semestre" in upper.lower():
        return "Campo semestre; revisar alcance."
    if "mu" in ptext or "matricula" in ptext:
        return "Evidencia proceso Matricula Unificada o matricula; no aplica automaticamente."
    return "Campo periodo/semestre detectado; requiere clasificacion por fuente."


def build_promedios_periodo_raw() -> pd.DataFrame:
    cols = ["CODCLI", "CODRAMO", "ANO", "PERIODO", "ESTADO", "DESCRIPCION_ESTADO"]
    df = read_xlsx(PROMEDIOS, "Hoja1", usecols=cols)
    for col in cols:
        df[col] = df[col].map(clean)
    rows = []
    years = " | ".join(sorted(v for v in df["ANO"].unique() if v))
    for periodo in sorted(df["PERIODO"].unique(), key=lambda x: (x == "", x)):
        subset = df[df["PERIODO"] == periodo]
        sub2025 = subset[subset["ANO"] == "2025"]
        estado_counts = (
            sub2025.groupby(["ESTADO", "DESCRIPCION_ESTADO"], dropna=False)
            .size()
            .reset_index(name="n")
            .sort_values(["ESTADO", "DESCRIPCION_ESTADO"])
        )
        estado_txt = "; ".join(
            f"{clean(r['ESTADO']) or '(blank)'}:{clean(r['DESCRIPCION_ESTADO']) or '(blank)'}={int(r['n'])}"
            for _, r in estado_counts.iterrows()
        )
        rows.append(
            {
                "Hoja": "Hoja1",
                "Campo PERIODO": "PERIODO",
                "Valores detectados": clean(periodo) or "(blank)",
                "Conteo de registros por PERIODO": len(subset),
                "Conteo registros 2025 por PERIODO": len(sub2025),
                "Conteo de CODCLI por PERIODO": sub2025["CODCLI"].nunique(),
                "Conteo de CODRAMO por PERIODO": sub2025["CODRAMO"].nunique(),
                "Conteo por ESTADO/DESCRIPCION_ESTADO si existe": estado_txt,
                "Años presentes": years,
                "Observación": "Dato observado B; no hay diccionario de significado de PERIODO raw en PROMEDIOS.",
            }
        )
    return pd.DataFrame(rows)


def keyify(df: pd.DataFrame, columns: Iterable[str]) -> pd.DataFrame:
    out = df.copy()
    for col in columns:
        if col in out.columns:
            out[col] = out[col].map(clean)
    return out


def build_cases_17() -> pd.DataFrame:
    cases = keyify(read_xlsx(H63_XLSX, "04_17_PERIODO_NO_1_2"), ["FILA_KEY", "FILA_5809"])
    evidence = keyify(read_xlsx(H63_XLSX, "06_EVIDENCIA_H55"), ["FILA_KEY"])
    merged = cases.merge(evidence, on="FILA_KEY", how="left", suffixes=("", "_H55"))
    rows = []
    for _, r in merged.iterrows():
        rows.append(
            {
                "FILA_5809": clean(r.get("FILA_5809")),
                "NUM_DOCUMENTO": clean(r.get("NUM_DOCUMENTO")),
                "CODIGO_UNICO": clean(r.get("CODIGO_UNICO")),
                "PLAN_ESTUDIOS": clean(r.get("PLAN_ESTUDIOS")),
                "CODCLI_LISTA": clean(r.get("CODCLI_LISTA")),
                "COLUMNA_AFECTADA": clean(r.get("COLUMNA_AFECTADA_16_19")),
                "VALOR_5809": clean(r.get("VALOR_ACTUAL_5809")),
                "VALOR_RECALCULADO_OBSERVADO": clean(r.get("VALOR_RECALCULADO_OBSERVADO")),
                "PERIODO raw observado": clean(r.get("PERIODOS_OBSERVADOS")),
                "Estado de gobernanza": "BLOQUEADO",
                "Dictamen por caso": "No aplicar criterio de PERIODO raw; requiere decision funcional/documental.",
            }
        )
    return pd.DataFrame(rows)


def build_detail_17() -> pd.DataFrame:
    detail = keyify(read_xlsx(H63_XLSX, "07_DETALLE_PERIODO_H55"), ["FILA_KEY", "FILA_5809"])
    cases = keyify(read_xlsx(H63_XLSX, "04_17_PERIODO_NO_1_2"), ["FILA_KEY", "FILA_5809"])[["FILA_KEY", "CODCLI_LISTA"]]
    merged = detail.merge(cases, on="FILA_KEY", how="left")
    rows = []
    for _, r in merged.iterrows():
        rows.append(
            {
                "FILA_5809": clean(r.get("FILA_5809")),
                "CODCLI_LISTA": clean(r.get("CODCLI_LISTA")),
                "CODCLI_NORM": clean(r.get("CODCLI_NORM")),
                "CODRAMO": clean(r.get("CODRAMO")),
                "PERIODO": clean(r.get("PERIODO")),
                "ESTADO": clean(r.get("ESTADO")),
                "DESCRIPCION_ESTADO": clean(r.get("DESCRIPCION_ESTADO")),
                "Año si existe": "NO_DISPONIBLE_EN_DETALLE_H55; el bloque viene de registros 2025 auditados",
                "Ramo duplicado SI/NO si existe": "MISMO_CODCLI=" + clean(r.get("DUPLICADO_MISMO_CODCLI")) + "; MULTIPLE_CODCLI=" + clean(r.get("DUPLICADO_MULTIPLE_CODCLI")),
                "Observación": "Detalle hito 55/63; evidencia tecnica de periodos observados.",
            }
        )
    return pd.DataFrame(rows)


def parse_int(value: Any) -> int:
    text = clean(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def scenario_prediction(group: pd.DataFrame, scenario: str) -> Dict[str, str]:
    work = group.copy()
    work["PERIODO_N"] = work["PERIODO"].map(clean)
    if scenario == "A":
        sem1, sem2, annual = {"1"}, {"2"}, {"1", "2"}
    elif scenario == "B":
        sem1, sem2, annual = {"1"}, {"2", "3"}, {"1", "2", "3"}
    elif scenario == "C":
        sem1, sem2, annual = {"1"}, {"2"}, {"1", "2"}
    elif scenario == "D":
        sem1, sem2, annual = {"1"}, {"2"}, {"1", "2", "3", "4", "5"}
    else:
        raise ValueError(f"Escenario no soportado: {scenario}")
    pvals = set(work["PERIODO_N"])
    annual_count = work[work["PERIODO_N"].isin(annual)]["CODRAMO"].map(clean).replace("", pd.NA).dropna().nunique()
    return {
        "CURSO_1ER_SEM": "SI" if pvals & sem1 else "NO",
        "CURSO_2DO_SEM": "SI" if pvals & sem2 else "NO",
        "UNIDADES_CURSADAS": str(annual_count),
    }


def build_scenario_detail() -> pd.DataFrame:
    ev = keyify(read_xlsx(H63_XLSX, "06_EVIDENCIA_H55"), ["FILA_KEY"])
    detail = keyify(read_xlsx(H63_XLSX, "07_DETALLE_PERIODO_H55"), ["FILA_KEY", "FILA_5809"])
    rows = []
    for _, case in ev.iterrows():
        fk = clean(case["FILA_KEY"])
        group = detail[detail["FILA_KEY"] == fk]
        current = {
            "CURSO_1ER_SEM": clean(case.get("CURSO_1ER_SEM_5809")),
            "CURSO_2DO_SEM": clean(case.get("CURSO_2DO_SEM_5809")),
            "UNIDADES_CURSADAS": clean(case.get("UNIDADES_CURSADAS_5809")),
        }
        for scenario in ["A", "B", "C", "D"]:
            pred = scenario_prediction(group, scenario)
            rows.append(
                {
                    "Escenario": scenario,
                    "FILA_5809": fk,
                    "PERIODOS_OBSERVADOS": clean(case.get("PERIODOS_OBSERVADOS")),
                    "5809_16": current["CURSO_1ER_SEM"],
                    "5809_17": current["CURSO_2DO_SEM"],
                    "5809_18": current["UNIDADES_CURSADAS"],
                    "SIM_16": pred["CURSO_1ER_SEM"],
                    "SIM_17": pred["CURSO_2DO_SEM"],
                    "SIM_18": pred["UNIDADES_CURSADAS"],
                    "CALZA_5809": "SI" if pred == current else "NO",
                    "COLUMNAS_CAMBIAN": " | ".join(
                        col
                        for col, label in [("16", "CURSO_1ER_SEM"), ("17", "CURSO_2DO_SEM"), ("18", "UNIDADES_CURSADAS")]
                        if pred[label] != current[label]
                    ),
                }
            )
        rows.append(
            {
                "Escenario": "E",
                "FILA_5809": fk,
                "PERIODOS_OBSERVADOS": clean(case.get("PERIODOS_OBSERVADOS")),
                "5809_16": current["CURSO_1ER_SEM"],
                "5809_17": current["CURSO_2DO_SEM"],
                "5809_18": current["UNIDADES_CURSADAS"],
                "SIM_16": current["CURSO_1ER_SEM"],
                "SIM_17": current["CURSO_2DO_SEM"],
                "SIM_18": current["UNIDADES_CURSADAS"],
                "CALZA_5809": "SI",
                "COLUMNAS_CAMBIAN": "",
            }
        )
    return pd.DataFrame(rows)


def build_scenarios() -> pd.DataFrame:
    detail = build_scenario_detail()
    desc = {
        "A": "Solo PERIODO 1 y 2.",
        "B": "1=SEM1, 2/3=SEM2.",
        "C": "1=SEM1, 2=SEM2, 3/4/5 bloqueados.",
        "D": "Todos los periodos como actividad anual, solo 1/2 para semestres.",
        "E": "Mantener 5809 sin cambio.",
    }
    risk = {
        "A": "ALTO: ignora actividad raw en 3/4/5 sin respaldo.",
        "B": "ALTO: aplica 2/3=SEM2 sin respaldo Avance; hito 64 lo bloquea.",
        "C": "MEDIO/ALTO: no resuelve casos con 3/4/5; mantiene bloqueo.",
        "D": "ALTO: cuenta 3/4/5 en anual sin significado funcional.",
        "E": "MEDIO: calza 5809, pero requiere decision funcional para cerrar sin cambio.",
    }
    level = {"A": "C/E", "B": "C/D/E", "C": "D/E", "D": "E", "E": "D"}
    verdict = {
        "A": "NO_PERMITIDO_COMO_REGLA_FINAL",
        "B": "NO_PERMITIDO_BLOQUEADO_H64",
        "C": "PENDIENTE",
        "D": "NO_PERMITIDO_SIN_VALIDACION",
        "E": "PENDIENTE_DECISION_FUNCIONAL",
    }
    rows = []
    for scenario, group in detail.groupby("Escenario"):
        calzan = int((group["CALZA_5809"] == "SI").sum())
        no_calzan = int((group["CALZA_5809"] == "NO").sum())
        changed_cols = []
        for col in ["16", "17", "18"]:
            changed_cols.append(f"{col}={int(group['COLUMNAS_CAMBIAN'].str.contains(col, regex=False).sum())}")
        rows.append(
            {
                "Escenario": f"Escenario {scenario}",
                "Descripcion": desc[scenario],
                "Cuantos casos calzan con 5809": calzan,
                "Cuantos no calzan": no_calzan,
                "Que columnas cambian": "; ".join(changed_cols),
                "Riesgo": risk[scenario],
                "Nivel de respaldo": level[scenario],
                "Dictamen": verdict[scenario],
            }
        )
    return pd.DataFrame(rows)


def build_decision_matrix(promedios: pd.DataFrame) -> pd.DataFrame:
    observed = {clean(r["Valores detectados"]): r for _, r in promedios.iterrows()}
    rows = []
    for periodo in ["1", "2", "3", "4", "5"]:
        obs = observed.get(periodo, {})
        if periodo == "1":
            treatment = "Candidato a primer semestre solo si validacion funcional confirma que PERIODO raw=1 equivale a primer semestre."
            c16, c17, c18 = "PENDIENTE", "NO", "PENDIENTE"
            state = "PENDIENTE"
            level = "B/C/D"
            official = "Instructivo define primer semestre, pero no define PERIODO raw=1."
            tech = "Hito 55 usa PERIODO=1 para auditar CURSO_1ER_SEM; codigo MU tambien infiere 1->1, no aplicable automatico."
        elif periodo == "2":
            treatment = "Candidato a segundo semestre solo si validacion funcional confirma que PERIODO raw=2 equivale a segundo semestre."
            c16, c17, c18 = "NO", "PENDIENTE", "PENDIENTE"
            state = "PENDIENTE"
            level = "B/C/D"
            official = "Instructivo define segundo semestre, pero no define PERIODO raw=2."
            tech = "Hito 55 usa PERIODO=2 para auditar CURSO_2DO_SEM; codigo MU infiere 2->2, no aplicable automatico."
        elif periodo == "3":
            treatment = "Bloqueado; no aplicar 2/3=SEM2 sin respaldo especifico Avance."
            c16, c17, c18 = "NO", "PENDIENTE/BLOQUEADO", "PENDIENTE/BLOQUEADO"
            state = "BLOQUEADO"
            level = "B/C/D/E"
            official = "Sin evidencia oficial en instructivo para PERIODO raw=3."
            tech = "Existe implementacion MU 3->2 y reporte MU; hito 64 bloquea reutilizacion automatica."
        elif periodo == "4":
            treatment = "Bloqueado; significado no confirmado."
            c16, c17, c18 = "NO", "NO", "PENDIENTE/BLOQUEADO"
            state = "BLOQUEADO"
            level = "B/E"
            official = "Sin evidencia oficial en instructivo para PERIODO raw=4."
            tech = "Observado en PROMEDIOS 2025 con baja frecuencia; sin diccionario ni mapeo."
        else:
            treatment = "Bloqueado; significado no confirmado."
            c16, c17, c18 = "NO", "NO", "PENDIENTE/BLOQUEADO"
            state = "BLOQUEADO"
            level = "B/E"
            official = "Sin evidencia oficial en instructivo para PERIODO raw=5."
            tech = "Observado en PROMEDIOS 2025 con baja frecuencia; sin diccionario ni mapeo."
        rows.append(
            {
                "PERIODO raw": periodo,
                "Evidencia oficial": official,
                "Evidencia dato observado": f"PROMEDIOS 2025 registros={obs.get('Conteo registros 2025 por PERIODO','')}; CODCLI={obs.get('Conteo de CODCLI por PERIODO','')}; CODRAMO={obs.get('Conteo de CODRAMO por PERIODO','')}",
                "Evidencia técnica": tech,
                "Decisión interna existente": f"{BLOQUEO_H64} {BLOQUEO_H66}",
                "Tratamiento candidato": treatment,
                "Aplicar a columna 16 SI/NO/PENDIENTE": c16,
                "Aplicar a columna 17 SI/NO/PENDIENTE": c17,
                "Aplicar a columna 18 SI/NO/PENDIENTE": c18,
                "Nivel respaldo": level,
                "Requiere validación funcional": "SI",
                "Estado": state,
            }
        )
    return pd.DataFrame(rows)


def build_mapeos_no_aplicar() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Mapeo": "PERIODO 2/3 = segundo semestre aplicado automaticamente",
                "Motivo": "Hito 64 lo clasifica pendiente/bloqueado; evidencia MU no gobierna Avance 5809.",
                "Fuente que lo bloquea": "Hito 64 / Hito 66",
                "Riesgo": "ALTO: altera columnas 17 y 18 sin respaldo Avance.",
                "Estado": "NO_APLICAR",
            },
            {
                "Mapeo": "PERIODO 3 = segundo semestre por codigo MU",
                "Motivo": "Codigo MU es implementacion tecnica de otro proceso; no regla oficial Avance.",
                "Fuente que lo bloquea": "Regla critica H64/H67",
                "Riesgo": "ALTO: reusa criterio transversal no confirmado.",
                "Estado": "NO_APLICAR",
            },
            {
                "Mapeo": "PERIODO 4 o 5 como actividad anual automaticamente",
                "Motivo": "No hay diccionario ni decision funcional sobre significado de 4/5.",
                "Fuente que lo bloquea": "Ausencia de respaldo oficial/diccionario PROMEDIOS",
                "Riesgo": "ALTO: suma unidades no gobernadas.",
                "Estado": "NO_APLICAR",
            },
            {
                "Mapeo": "PERIODO 1=SEM1 y PERIODO 2=SEM2 como regla definitiva sin validacion",
                "Motivo": "Uso tecnico trazable existe, pero no hay diccionario raw ni definicion oficial del campo PERIODO.",
                "Fuente que lo bloquea": "Hito 64 / objetivo hito 67",
                "Riesgo": "MEDIO: puede ser correcto, pero requiere formalizacion.",
                "Estado": "NO_APLICAR_TODAVIA",
            },
            {
                "Mapeo": "Reglas de Matricula Unificada para Avance Curricular",
                "Motivo": "Procesos y campos distintos; MU no reemplaza instructivo Avance.",
                "Fuente que lo bloquea": "Hito 64",
                "Riesgo": "ALTO",
                "Estado": "NO_APLICAR",
            },
        ]
    )


def build_recomendacion() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Pregunta": "Que puede usarse ahora",
                "Respuesta": "Inventario, conteos PROMEDIOS, detalle H55/H63 y bloqueo H64/H66 como evidencia documental.",
                "Efecto sobre 17 casos": "Ninguno; no corrige ni destraba automaticamente.",
                "Estado": "PERMITIDO_COMO_EVIDENCIA",
            },
            {
                "Pregunta": "Que no puede usarse",
                "Respuesta": "PERIODO 2/3=SEM2, PERIODO 4/5 anual, ni reglas MU como regla Avance.",
                "Efecto sobre 17 casos": "Mantiene bloqueo.",
                "Estado": "BLOQUEADO",
            },
            {
                "Pregunta": "Que requiere validacion",
                "Respuesta": "Significado institucional de PERIODO raw 1,2,3,4,5 en PROMEDIOS.",
                "Efecto sobre 17 casos": "Podria moverlos a decision funcional o mantenerlos bloqueados.",
                "Estado": "PENDIENTE",
            },
            {
                "Pregunta": "Decision funcional requerida",
                "Respuesta": "Confirmar si 3 corresponde a segundo semestre, periodo especial, recuperacion, verano u otro; y tratamiento de 4/5.",
                "Efecto sobre 17 casos": "Permite definir si mantener 5809 sin cambio o preparar correccion futura gobernada.",
                "Estado": "REQUERIDA",
            },
            {
                "Pregunta": "Se puede destrabar alguno sin correccion",
                "Respuesta": "Solo con decision funcional documentada. En simulacion, mantener 5809 calza 17/17, pero no constituye regla por si solo.",
                "Efecto sobre 17 casos": "Podria cerrar sin cambio si area funcional valida.",
                "Estado": "PENDIENTE_DECISION_FUNCIONAL",
            },
            {
                "Pregunta": "20-21",
                "Respuesta": "Fuera de alcance; no evaluado y no recalculado.",
                "Efecto sobre 17 casos": "No aplica.",
                "Estado": "SEPARADO",
            },
        ]
    )


def build_borrador_text() -> str:
    return f"""# Borrador decision funcional PERIODO raw 5809

Proceso: {PROCESO}
Subproyecto: {SUBPROYECTO}
Ano proceso: {ANIO_PROCESO}
Ano referencia datos: {ANIO_REFERENCIA}
Declaracion carga: {DECLARACION_CARGA}

Estimadas/os,

Se solicita validacion funcional especifica sobre el significado institucional del campo PERIODO raw observado en PROMEDIOS para su eventual uso en el archivo 5809 Matricula Avance Curricular, columnas 16, 17 y 18.

Esta solicitud no implica carga, no modifica fuentes originales, no corrige datos, no genera SIES_READY y no recalcula columnas 20-21.

## Preguntas a validar

1. Confirmar el significado institucional de PERIODO raw 1, 2, 3, 4 y 5 en PROMEDIOS.
2. Confirmar si PERIODO 1 corresponde a primer semestre para Avance Curricular 5809.
3. Confirmar si PERIODO 2 corresponde a segundo semestre para Avance Curricular 5809.
4. Confirmar si PERIODO 3 corresponde a segundo semestre, periodo especial, recuperacion, verano u otro significado.
5. Confirmar si PERIODO 4 y PERIODO 5 deben contarse en UNIDADES_CURSADAS anual 2025.
6. Confirmar si PERIODO 4 y PERIODO 5 deben afectar CURSO_1ER_SEM o CURSO_2DO_SEM.
7. Confirmar si los 17 casos PERIODO_NO_1_2 deben mantenerse como 5809 o corregirse en un hito posterior.
8. Confirmar que la decision aplica solo a Avance Curricular 5809 columnas 16-18 y no a Matricula Unificada ni a columnas 20-21.

## Evidencia resumida

- PROMEDIOS 2025 contiene PERIODO 1,2,3,4 y 5.
- Hito 55 detecto 17 casos con causa principal PERIODO_NO_1_2.
- Hito 64 bloqueo el uso automatico de PERIODO 2/3=segundo semestre.
- Hito 66 mantuvo bloqueados los 17 casos hasta gobernanza especifica.
- La simulacion documental del hito 67 muestra que mantener 5809 calza 17/17, pero requiere decision funcional para ser cierre valido.

Resultado esperado:

- Definir tratamiento por PERIODO raw 1,2,3,4,5.
- Autorizar o bloquear su uso en columnas 16, 17 y 18.
- Indicar si los 17 casos se mantienen sin cambio o deben pasar a correccion posterior gobernada.
"""


def build_borrador_sheet(text: str) -> pd.DataFrame:
    return pd.DataFrame([{"Orden": i, "Texto": part.strip()} for i, part in enumerate(text.split("\n\n"), start=1)])


def build_bloqueos() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"Bloqueo/Pendiente": "PERIODO 3", "Descripcion": "No confirmado como segundo semestre ni periodo especial.", "Fuente": "PROMEDIOS B + H64/H66 D/E", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "PERIODO 4", "Descripcion": "Observado en PROMEDIOS 2025; significado no confirmado.", "Fuente": "PROMEDIOS B", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "PERIODO 5", "Descripcion": "Observado en PROMEDIOS 2025; significado no confirmado.", "Fuente": "PROMEDIOS B", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "Falta de diccionario", "Descripcion": "No se detecto diccionario de PERIODO raw en PROMEDIOS.", "Fuente": "PROMEDIOS", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "Falta de respaldo oficial", "Descripcion": "Instructivo no define PERIODO raw 1..5.", "Fuente": "Instructivo A", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "Contradiccion potencial MU/Avance", "Descripcion": "MU usa 3->2 tecnicamente; H64 impide reuso automatico para Avance.", "Fuente": "Codigo MU C + H64", "Estado": "BLOQUEADO"},
            {"Bloqueo/Pendiente": "17 casos PERIODO_NO_1_2", "Descripcion": "No se destraban sin validacion funcional/documental.", "Fuente": "H55/H63/H66", "Estado": "BLOQUEADO"},
        ]
    )


def build_fuentes(output_paths: Dict[str, Path], desktop_dir: Path) -> pd.DataFrame:
    source_paths = [
        ("A", "Instructivo oficial Avance", INSTRUCTIVO, "Regla oficial de conceptos semestre/unidades cursadas."),
        ("B", "PROMEDIOS actualizado", PROMEDIOS, "Dato observado PERIODO raw."),
        ("C/D", "Hito 55 carpeta", H55_DIR, "Auditoria diferencia 16-18 multiple."),
        ("C/D", "Hito 55 Excel", H55_XLSX, "Detalle ramo/periodo y causas."),
        ("C", "Hito 55 script", H55_SCRIPT, "Implementacion tecnica que detecta PERIODO_NO_1_2."),
        ("C/D", "Hito 63 carpeta", H63_DIR, "Revision 38 restante."),
        ("C/D", "Hito 63 Excel", H63_XLSX, "17 casos y detalle H55."),
        ("D/E", "Hito 64 carpeta", H64_DIR, "Exploracion transversal y bloqueo PERIODO raw."),
        ("D/E", "Hito 64 Excel", H64_XLSX, "Periodo academico y bloqueos."),
        ("D", "Hito 66 carpeta", H66_DIR, "Correccion dictamen y prompt PERIODO raw."),
        ("D", "Hito 66 Excel", H66_XLSX, "Bloque 17 y prompt."),
        ("C", "Control FOR_ING_ACT", CONTROL_FOR_ING, "Control MU/ingreso; no regla Avance 5809."),
        ("C", "Control CODCARPR", CONTROL_CODCARPR, "Control local; sin PERIODO raw."),
        ("C", "Control campos ingreso", CONTROL_CAMPOS_ING, "Control SEM_ING/PERIODOINGRESO; no PERIODO raw ramos."),
        ("C", "Codigo gobernanza MU", MU_CODE, "Implementacion MU; no aplicar automaticamente."),
        ("C/D", "Auditoria MU", MU_AUDIT, "Reporte MU; no aplicar automaticamente."),
    ]
    rows = []
    for idx, (level, tipo, path, usage) in enumerate(source_paths, start=1):
        rows.append(
            {
                "ID": f"F{idx:03d}",
                "Tipo": tipo,
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() and path.is_file() else "",
                "Nivel respaldo": level,
                "Uso": usage,
            }
        )
    for key, path in output_paths.items():
        rows.append(
            {
                "ID": f"OUT_{key.upper()}",
                "Tipo": "Producto hito 67",
                "Ruta": str(path),
                "Existe": "SI" if path.exists() else "NO",
                "Hash SHA256": sha256_file(path) if path.exists() and path.is_file() and key != "manifest" else "",
                "Nivel respaldo": "DERIVADO",
                "Uso": "Producto documental; no fuente normativa.",
            }
        )
    rows.append(
        {
            "ID": "OUT_DESKTOP",
            "Tipo": "Copia escritorio",
            "Ruta": str(desktop_dir),
            "Existe": "SI" if desktop_dir.exists() else "NO",
            "Hash SHA256": "",
            "Nivel respaldo": "DERIVADO",
            "Uso": "Copia de productos al Escritorio.",
        }
    )
    return pd.DataFrame(rows)


def build_dictamen(promedios: pd.DataFrame, matrix: pd.DataFrame, cases17: pd.DataFrame) -> pd.DataFrame:
    periods = [p for p in promedios["Valores detectados"].tolist() if p != "(blank)"]
    governed = int((matrix["Estado"] == "GOBERNADO").sum())
    blocked = int((matrix["Estado"] == "BLOQUEADO").sum())
    return pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Año proceso": ANIO_PROCESO,
                "Año referencia": ANIO_REFERENCIA,
                "Estado": ESTADO,
                "Declaración carga": DECLARACION_CARGA,
                "Fuentes originales modificadas": "NO",
                "Correcciones aplicadas": "NO",
                "SIES_READY generado": "NO",
                "20-21 evaluado": "NO",
                "Total casos PERIODO_NO_1_2 revisados": len(cases17),
                "Cantidad de períodos raw detectados": len(periods),
                "Períodos raw detectados": " | ".join(periods),
                "Períodos gobernados": governed,
                "Períodos bloqueados": blocked,
                "Dictamen global": DICTAMEN_GLOBAL,
            }
        ]
    )


def build_manifest(timestamp: str, output_dir: Path, desktop_dir: Path, output_paths: Dict[str, Path], sheets: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    return {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": timestamp,
        "estado": ESTADO,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": "NO",
        "correcciones_aplicadas": "NO",
        "sies_ready_generado": "NO",
        "archivo_carga_generado": "NO",
        "recalculo_20_21": "NO",
        "reglas_periodo_raw_aplicadas": "NO",
        "reglas_matricula_unificada_aplicadas": "NO",
        "total_casos_periodo_no_1_2": 17,
        "carpeta_salida": str(output_dir),
        "carpeta_escritorio": str(desktop_dir),
        "productos": {key: str(path) for key, path in output_paths.items()},
        "hojas_excel": {logical: {"hoja_excel": SHEETS.get(logical, logical)[:31], "filas": len(df)} for logical, df in sheets.items()},
        "hashes_salida": {key: sha256_file(path) for key, path in output_paths.items() if path.exists() and path.is_file() and key != "manifest"},
    }


def manifest_legible(manifest: Dict[str, Any]) -> pd.DataFrame:
    rows = []
    for key, value in manifest.items():
        rows.append({"Clave": key, "Valor": json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value})
    return pd.DataFrame(rows)


def excel_safe_frame(df: pd.DataFrame) -> pd.DataFrame:
    def safe(value: Any) -> Any:
        if isinstance(value, str) and len(value) > 32000:
            return value[:31980] + " [...TRUNCADO_PARA_EXCEL]"
        return value

    return df.map(safe)


def write_excel(path: Path, sheets: Dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for logical, df in sheets.items():
            sheet_name = SHEETS.get(logical, logical)[:31]
            excel_safe_frame(df).to_excel(writer, sheet_name=sheet_name, index=False)
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx, column_cells in enumerate(ws.columns, start=1):
            values = [display(cell.value, 80) for cell in column_cells[:120]]
            width = min(max([len(v) for v in values] + [10]) + 2, 60)
            ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(path)


def build_markdown(
    dictamen: pd.DataFrame,
    extractos: pd.DataFrame,
    promedios: pd.DataFrame,
    cases17: pd.DataFrame,
    scenarios: pd.DataFrame,
    matrix: pd.DataFrame,
    mapeos_no: pd.DataFrame,
    recomendacion: pd.DataFrame,
    bloqueos: pd.DataFrame,
    output_paths: Dict[str, Path],
    desktop_dir: Path,
) -> str:
    return f"""# Gobernanza PERIODO raw 5809

## Estado

- Proceso: {PROCESO}
- Subproyecto: {SUBPROYECTO}
- Ano proceso: {ANIO_PROCESO}
- Ano referencia datos: {ANIO_REFERENCIA}
- Estado: {ESTADO}
- Declaracion carga: {DECLARACION_CARGA}
- Fuentes originales modificadas: NO
- Correcciones aplicadas: NO
- SIES_READY generado: NO
- Archivo de carga generado: NO
- 20-21 evaluado: NO

## Que se reviso

Se revisaron el instructivo oficial Avance, PROMEDIOS actualizado, hito 55, hito 63, hito 64, hito 66 y archivos locales de control/codigo/reporte con referencias a PERIODO, PERIODOMATRICULA, SEMESTRE o SEM_INGRESO.

## Instructivo

El instructivo gobierna los conceptos de presencia en primer y segundo semestre y unidades cursadas durante el ano academico 2025. No define el campo raw PROMEDIOS `PERIODO` ni entrega diccionario para valores 1,2,3,4,5.

{markdown_table(extractos[["Texto exacto", "Campo afectado", "Linea/seccion", "Observación"]], max_rows=14)}

## PERIODO raw en PROMEDIOS

PROMEDIOS Hoja1 contiene `PERIODO` con valores 1,2,3,4,5 y blancos. Para 2025 se observan registros en 1,2,3,4 y 5. Esto es dato observado B, no regla oficial.

{markdown_table(promedios[["Valores detectados", "Conteo registros 2025 por PERIODO", "Conteo de CODCLI por PERIODO", "Conteo de CODRAMO por PERIODO"]])}

## Evidencia por periodo

{markdown_table(matrix[["PERIODO raw", "Evidencia oficial", "Evidencia dato observado", "Evidencia técnica", "Tratamiento candidato", "Estado"]])}

## Los 17 casos

Los 17 casos vienen de hito 63 y mantienen estado BLOQUEADO. No se aplica criterio 1=SEM1 ni 2/3=SEM2.

{markdown_table(cases17[["FILA_5809", "CODIGO_UNICO", "CODCLI_LISTA", "PERIODO raw observado", "Estado de gobernanza"]], max_rows=17)}

## Escenarios

Los escenarios se simularon sin aplicar cambios. El escenario E, mantener 5809, calza 17/17; eso no basta como regla, pero sirve como alternativa a validar funcionalmente. El escenario B, 1=SEM1 y 2/3=SEM2, solo calza 1/17 y esta bloqueado por hito 64.

{markdown_table(scenarios)}

## Por que no se puede aplicar todavia

{markdown_table(mapeos_no)}

## Recomendacion funcional

{markdown_table(recomendacion)}

## Bloqueos

{markdown_table(bloqueos)}

## Por que sigue NO_LISTO_PARA_CARGA

El hito no aplica reglas, no corrige datos, no genera carga, no genera SIES_READY y deja los 17 casos bloqueados hasta decision funcional/documental. Las columnas 20-21 permanecen separadas y no evaluadas.

## Archivos

- Excel: {output_paths['excel']}
- Informe: {output_paths['informe']}
- Manifest: {output_paths['manifest']}
- Script: {output_paths['script']}
- Borrador decision: {output_paths['borrador']}
- Carpeta escritorio: {desktop_dir}
"""


def copy_to_desktop(output_dir: Path, desktop_dir: Path) -> None:
    desktop_dir.parent.mkdir(parents=True, exist_ok=True)
    if desktop_dir.exists():
        shutil.rmtree(desktop_dir)
    shutil.copytree(output_dir, desktop_dir, ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))


def print_console(
    dictamen: pd.DataFrame,
    promedios: pd.DataFrame,
    scenarios: pd.DataFrame,
    matrix: pd.DataFrame,
    cases17: pd.DataFrame,
    bloqueos: pd.DataFrame,
    output_paths: Dict[str, Path],
    desktop_dir: Path,
) -> None:
    print("GOBERNANZA PERIODO RAW 5809 — GENERADA")
    print("Fuentes originales modificadas: NO")
    print("Correcciones aplicadas: NO")
    print("SIES_READY generado: NO")
    print("Archivo de carga generado: NO")
    print("20-21 evaluado: NO")
    print("Declaración carga: NO_LISTO_PARA_CARGA")
    print()
    print("DICTAMEN GLOBAL")
    print(dictamen.to_string(index=False))
    print()
    print("PERIODOS RAW DETECTADOS")
    print(promedios[["Valores detectados", "Conteo registros 2025 por PERIODO", "Conteo de CODCLI por PERIODO", "Conteo de CODRAMO por PERIODO"]].to_string(index=False))
    print()
    print("ESCENARIOS DE MAPEO")
    print(scenarios.to_string(index=False))
    print()
    print("MATRIZ DECISION PERIODO")
    print(matrix[["PERIODO raw", "Tratamiento candidato", "Nivel respaldo", "Requiere validación funcional", "Estado"]].to_string(index=False))
    print()
    print("17 CASOS")
    resumen17 = pd.DataFrame(
        [
            {
                "Casos": len(cases17),
                "Periodos observados": cases17["PERIODO raw observado"].value_counts().to_dict(),
                "Estado": cases17["Estado de gobernanza"].value_counts().to_dict(),
            }
        ]
    )
    print(resumen17.to_string(index=False))
    print()
    print("BLOQUEOS")
    print(bloqueos.to_string(index=False))
    print()
    print("ARCHIVOS")
    print(f"Excel: {output_paths['excel']}")
    print(f"Informe: {output_paths['informe']}")
    print(f"Manifest: {output_paths['manifest']}")
    print(f"Script: {output_paths['script']}")
    print(f"Borrador decisión: {output_paths['borrador']}")
    print(f"Carpeta escritorio: {desktop_dir}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Gobernanza PERIODO raw 5809")
    parser.add_argument("--timestamp", default=None, help="Timestamp YYYYMMDD_HHMMSS")
    args = parser.parse_args()
    timestamp = args.timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")

    output_dir = BASE_OUT / f"GOBERNANZA_PERIODO_RAW_5809_{timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    desktop_dir = Path.home() / "Desktop" / f"AVANCE_CURRICULAR_2026_GOBERNANZA_PERIODO_RAW_5809_{timestamp}"
    output_paths = {key: output_dir / filename for key, filename in OUTPUT_FILES.items()}
    current_script = Path(__file__).resolve()
    if current_script != output_paths["script"].resolve():
        shutil.copy2(current_script, output_paths["script"])

    sources = build_sources(output_dir)
    extractos = build_extractos_instructivo()
    campos = build_campos_periodo_detectados(sources)
    promedios = build_promedios_periodo_raw()
    cases17 = build_cases_17()
    detail17 = build_detail_17()
    scenarios = build_scenarios()
    matrix = build_decision_matrix(promedios)
    mapeos_no = build_mapeos_no_aplicar()
    recomendacion = build_recomendacion()
    borrador_text = build_borrador_text()
    borrador_sheet = build_borrador_sheet(borrador_text)
    bloqueos = build_bloqueos()
    dictamen = build_dictamen(promedios, matrix, cases17)

    sheets: Dict[str, pd.DataFrame] = OrderedDict(
        [
            ("00_DICTAMEN_GLOBAL", dictamen),
            ("01_FUENTES_REVISADAS", sources),
            ("02_EXTRACTOS_INSTRUCTIVO", extractos),
            ("03_CAMPOS_PERIODO_DETECTADOS", campos),
            ("04_PROMEDIOS_PERIODO_RAW", promedios),
            ("05_17_CASOS_PERIODO_NO_1_2", cases17),
            ("06_DETALLE_RAMO_PERIODO_17", detail17),
            ("07_ESCENARIOS_DE_MAPEO", scenarios),
            ("08_MATRIZ_DECISION_PERIODO", matrix),
            ("09_MAPEOS_NO_APLICAR", mapeos_no),
            ("10_RECOMENDACION_FUNCIONAL", recomendacion),
            ("11_BORRADOR_DECISION_FUNCIONAL", borrador_sheet),
            ("12_BLOQUEOS_Y_PENDIENTES", bloqueos),
        ]
    )

    output_paths["borrador"].write_text(borrador_text, encoding="utf-8")
    output_paths["informe"].write_text(
        build_markdown(dictamen, extractos, promedios, cases17, scenarios, matrix, mapeos_no, recomendacion, bloqueos, output_paths, desktop_dir),
        encoding="utf-8",
    )
    sheets["13_FUENTES"] = build_fuentes(output_paths, desktop_dir)
    manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets)
    sheets["14_MANIFEST_LEGIBLE"] = manifest_legible(manifest)
    write_excel(output_paths["excel"], sheets)
    sheets["13_FUENTES"] = build_fuentes(output_paths, desktop_dir)
    manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets)
    sheets["14_MANIFEST_LEGIBLE"] = manifest_legible(manifest)
    write_excel(output_paths["excel"], sheets)
    output_paths["manifest"].write_text(json.dumps(build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets), ensure_ascii=False, indent=2), encoding="utf-8")
    sheets["13_FUENTES"] = build_fuentes(output_paths, desktop_dir)
    manifest = build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets)
    sheets["14_MANIFEST_LEGIBLE"] = manifest_legible(manifest)
    write_excel(output_paths["excel"], sheets)
    output_paths["manifest"].write_text(json.dumps(build_manifest(timestamp, output_dir, desktop_dir, output_paths, sheets), ensure_ascii=False, indent=2), encoding="utf-8")

    copy_to_desktop(output_dir, desktop_dir)
    print_console(dictamen, promedios, scenarios, matrix, cases17, bloqueos, output_paths, desktop_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
