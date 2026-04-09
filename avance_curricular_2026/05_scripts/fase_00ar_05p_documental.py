#!/usr/bin/env python3
"""Rectificación y diseño documental Avance Curricular SIES 2026.

Ejecuta Fase 00A-R y genera artefactos documentales Fase 04-05P.
No completa datos reales, no genera PES y no modifica fuentes originales
fuera del subproyecto.
"""

from __future__ import annotations

import csv
import hashlib
import importlib.util
import json
import os
import re
import stat
import subprocess
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook


REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
ROOT = REPO / "avance_curricular_2026"
FREEZE = ROOT / "00_fuentes_congeladas" / "CARGA_CONGELADA_20260626_005826"
ORIG = FREEZE / "originales"
MAN = FREEZE / "manifiestos"
DOWNLOADS_PDF = Path("/Users/alexi/Downloads/Instructivo_Avance Curricular SIES - 2026.pdf")
TXT_NAME = "Instructivo_Avance Curricular SIES - 2026.txt"
PDF_NAME = "Instructivo_Avance Curricular SIES - 2026.pdf"
PROMEDIOS_NAME = "PROMEDIOSDEALUMNOS_7804.xlsx"
PROCESS = "Avance Curricular SIES 2026"
YEAR_PROCESS = "2026"
YEAR_DATA = "2025"
PERIOD = "Año académico 2025: 1er semestre, 2do semestre, avance anual y avance acumulado al cierre 2025"
TZ = ZoneInfo("America/Santiago")
NOW = datetime.now(TZ)
DECISION_ID = "DI-AC-2026-001"
SCRIPT_BASE = ROOT / "05_scripts" / "fase_00a_03_inicial.py"


def load_base_module() -> Any:
    spec = importlib.util.spec_from_file_location("fase_00a_03_inicial", SCRIPT_BASE)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


BASE = load_base_module()
CARRERAS_FIELDS = BASE.CARRERAS_FIELDS
MATRICULA_FIELDS = BASE.MATRICULA_FIELDS
FIELD_DEFINITIONS = BASE.FIELD_DEFINITIONS
CARRERAS_MODIFIABLE = BASE.CARRERAS_MODIFIABLE
MATRICULA_MODIFIABLE = BASE.MATRICULA_MODIFIABLE
PERSONAL_FIELD_NAMES = set(BASE.PERSONAL_FIELD_NAMES) | {
    "MAIL",
    "MAIL_INST",
    "FECHANACIMIENTO",
    "DIRECCIONPROCEDENCIA",
    "DIRECCIONACTUAL",
    "FONOACTUAL",
    "FONOPROCEDENCIA",
    "RUTRESPONSABLEFINANCIERO",
    "NOMBRERESPONSABLEFINANCIERO",
}


def run(cmd: list[str]) -> str:
    proc = subprocess.run(
        cmd,
        cwd=str(REPO),
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
    )
    return proc.stdout.strip()


def nfc(value: str) -> str:
    return unicodedata.normalize("NFC", value)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_dirs() -> None:
    for path in [
        ROOT / "01_documentacion",
        ROOT / "03_gobernanza_columnas" / "carreras",
        ROOT / "03_gobernanza_columnas" / "matricula",
        ROOT / "04_configuracion",
        ROOT / "07_control",
        ROOT / "08_auditorias",
        ROOT / "09_pendientes",
        ROOT / "05_scripts",
    ]:
        path.mkdir(parents=True, exist_ok=True)


def stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "SI" if value else "NO"
    if isinstance(value, (list, tuple, set)):
        return " | ".join(stringify(v) for v in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def write_tsv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({field: stringify(row.get(field, "")) for field in fields})


def read_tsv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh, delimiter="\t")
        return list(reader.fieldnames or []), list(reader)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2, sort_keys=True)
        fh.write("\n")


def safe_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    asciiish = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return re.sub(r"[^A-Za-z0-9]+", "_", asciiish).strip("_").upper() or "SIN_NOMBRE"


def mask_inventory(value: str) -> str:
    value = re.sub(r"(?i)(rut[_ -]?)[0-9][0-9._ -]{5,15}[0-9Kk]?", r"\1<ID_MASKED>", value)
    value = re.sub(r"(?i)(run[_ -]?)[0-9][0-9._ -]{5,15}[0-9Kk]?", r"\1<ID_MASKED>", value)
    return value


def classify_path(path: str) -> tuple[str, str, str]:
    p = Path(path)
    s = path
    if s.startswith("avance_curricular_2026/00_fuentes_congeladas"):
        personal = "SI" if any(token in s for token in ["MATRICULA", "PROMEDIOS", "PROMEDIOSDEALUMNOS"]) else "NO_DETERMINADO"
        return "DATO_PERSONAL_LOCAL" if personal == "SI" else "GENERADO_FASE_00A_03", personal, "IGNORAR_GIT"
    if s.startswith("avance_curricular_2026/03_gobernanza_columnas") or s.startswith("avance_curricular_2026/04_configuracion"):
        return "GENERADO_FASE_04", "NO", "CANDIDATO_COMMIT_POSTERIOR"
    if s.startswith("avance_curricular_2026/08_auditorias") or "DECISION_INTERNA" in s:
        return "GENERADO_RECTIFICACION", "NO", "CANDIDATO_COMMIT_POSTERIOR"
    if s.startswith("avance_curricular_2026/09_pendientes"):
        return "GENERADO_FASE_04", "NO", "CANDIDATO_COMMIT_POSTERIOR"
    if s.startswith("avance_curricular_2026/07_control") or s.startswith("avance_curricular_2026/01_documentacion"):
        return "GENERADO_FASE_04" if p.name not in {"00_DIAGNOSTICO_REPOSITORIO.md", "01_MATRIZ_NORMATIVA.md", "02_REPORTE_PRECARGAS.md"} else "GENERADO_FASE_00A_03", "NO", "CANDIDATO_COMMIT_POSTERIOR"
    if s.startswith("avance_curricular_2026/05_scripts"):
        return "GENERADO_FASE_04" if "00ar_05p" in s else "GENERADO_FASE_00A_03", "NO", "CANDIDATO_COMMIT_POSTERIOR"
    if s.startswith(("estudiantes_extranjeros_2026/", "personal_academico_2026/", "cned/", "indices_2025/")):
        return "PREEXISTENTE_OTRO_PROCESO", "NO_DETERMINADO", "NO_TOCAR"
    if s.startswith(("archive/", "backups/", "resultados/", "control/", "scripts/", "tests/")):
        return "PREEXISTENTE_OTRO_PROCESO", "NO_DETERMINADO", "REVISAR_MANUALMENTE"
    return "DESCONOCIDO", "NO_DETERMINADO", "REVISAR_MANUALMENTE"


def git_check_ignore(path: str) -> str:
    proc = subprocess.run(
        ["git", "check-ignore", "-q", "--", path],
        cwd=str(REPO),
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=False,
    )
    return "SI" if proc.returncode == 0 else "NO"


def phase_00ar() -> dict[str, Any]:
    pdf_candidates = sorted(FREEZE.rglob(PDF_NAME))
    if len(pdf_candidates) != 1:
        raise RuntimeError(f"PDF congelado esperado no encontrado de forma única: {pdf_candidates}")
    pdf_path = pdf_candidates[0]
    if pdf_path.suffix.lower() != ".pdf":
        raise RuntimeError("La ruta localizada no tiene extensión .pdf")
    if nfc(pdf_path.name) != nfc(PDF_NAME):
        raise RuntimeError("El nombre normalizado del PDF no corresponde al instructivo")
    if pdf_path.is_symlink():
        raise RuntimeError("La ruta del PDF es un enlace simbólico; eliminación bloqueada")
    resolved = pdf_path.resolve()
    freeze_resolved = FREEZE.resolve()
    if freeze_resolved not in resolved.parents:
        raise RuntimeError("La ruta resuelta del PDF escapa de la carga congelada")

    pdf_hash = sha256(pdf_path)
    pdf_size = pdf_path.stat().st_size
    txt_path = ORIG / TXT_NAME
    carreras_path = ORIG / "5810_Precarga Carreras Avance Curricular 20268.csv"
    matricula_paths = list(ORIG.glob("5809_Precarga Matri*cula Avance Curricular 2026.csv"))
    if len(matricula_paths) != 1:
        raise RuntimeError("No se encontró de forma única la precarga de matrícula congelada")
    matricula_path = matricula_paths[0]
    promedios_path = ORIG / PROMEDIOS_NAME
    before = {
        "txt": sha256(txt_path),
        "carreras": sha256(carreras_path),
        "matricula": sha256(matricula_path),
        "promedios": sha256(promedios_path),
    }

    if not DOWNLOADS_PDF.exists():
        raise RuntimeError("El PDF original de Descargas no existe; eliminación bloqueada")

    pdf_path.chmod(stat.S_IRUSR | stat.S_IWUSR | stat.S_IRGRP | stat.S_IROTH)
    pdf_path.unlink()

    comparison_path = FREEZE / "evidencia" / ("COMPARACION_INSTRUCTIVO_TXT_" + "PDF.json")
    comparison_removed = "NO"
    if comparison_path.exists():
        comparison_path.unlink()
        comparison_removed = "SI"

    retiro_time = NOW.isoformat()
    manifest_path = MAN / "MANIFIESTO_CARGA_CONGELADA.tsv"
    fields, rows = read_tsv(manifest_path)
    for field in ["ACTIVO_EN_CARGA", "FECHA_RETIRO", "MOTIVO_RETIRO", "DECISION_INTERNA_ID"]:
        if field not in fields:
            fields.append(field)
    for row in rows:
        row.setdefault("ACTIVO_EN_CARGA", "SI")
        row.setdefault("FECHA_RETIRO", "NO_APLICA")
        row.setdefault("MOTIVO_RETIRO", "NO_APLICA")
        row.setdefault("DECISION_INTERNA_ID", "NO_APLICA")
        if row.get("ID_FUENTE") == "FUENTE_INSTRUCTIVO_PDF":
            row["ESTADO_VALIDACION"] = "RETIRADO_POR_DECISION_INTERNA"
            row["HASH_COINCIDE"] = "NO_APLICA"
            row["ACTIVO_EN_CARGA"] = "NO"
            row["FECHA_RETIRO"] = retiro_time
            row["MOTIVO_RETIRO"] = "PDF_NO_REQUERIDO_FUENTE_TXT_UNICA"
            row["DECISION_INTERNA_ID"] = DECISION_ID
            row["OBSERVACIONES"] = (
                "Copia retirada del proyecto por decisión interna. El TXT se establece como fuente normativa única. "
                "El original permanece en Descargas y no fue modificado."
            )
        elif row.get("ESTADO_VALIDACION") == "CONGELADO_VALIDADO":
            row["ACTIVO_EN_CARGA"] = "SI"
    write_tsv(manifest_path, rows, fields)

    inv_path = ROOT / "07_control" / "INVENTARIO_FUENTES.tsv"
    inv_fields, inv_rows = read_tsv(inv_path)
    for field in ["NIVEL_RESPALDO", "ACTIVO", "FECHA_RETIRO", "DECISION_INTERNA_ID"]:
        if field not in inv_fields:
            inv_fields.append(field)
    for row in inv_rows:
        row.setdefault("NIVEL_RESPALDO", "NO_DETERMINADO")
        row.setdefault("ACTIVO", "SI" if row.get("ID_FUENTE") != "EXCLUIDO" else "NO")
        row.setdefault("FECHA_RETIRO", "NO_APLICA")
        row.setdefault("DECISION_INTERNA_ID", "NO_APLICA")
        if row.get("ID_FUENTE") == "FUENTE_INSTRUCTIVO_PDF":
            row["CLASIFICACION"] = "EVIDENCIA_RETIRADA"
            row["NIVEL_RESPALDO"] = "NO_APLICA_EN_PROCESO_ACTIVO"
            row["ESTADO"] = "RETIRADO_POR_DECISION_INTERNA"
            row["ACTIVO"] = "NO"
            row["FECHA_RETIRO"] = retiro_time
            row["DECISION_INTERNA_ID"] = DECISION_ID
            row["APTO_PARA_GIT"] = "NO"
            row["OBSERVACIONES"] = "El TXT queda como fuente normativa única; PDF retirado del subproyecto."
        elif row.get("ID_FUENTE") != "EXCLUIDO":
            row["ACTIVO"] = "SI"
    write_tsv(inv_path, inv_rows, inv_fields)

    after = {
        "txt": sha256(txt_path),
        "carreras": sha256(carreras_path),
        "matricula": sha256(matricula_path),
        "promedios": sha256(promedios_path),
    }
    if before != after:
        raise RuntimeError(f"Hash de fuentes restantes cambió: before={before} after={after}")

    audit_rows = [
        {
            "ID_ACCION": "RETIRO-PDF-AC-2026-001",
            "FECHA_HORA": retiro_time,
            "ZONA_HORARIA": "America/Santiago",
            "PROCESO": PROCESS,
            "SUBPROCESO": "Proceso Avance Curricular SIES 2026",
            "DECISION_INTERNA_ID": DECISION_ID,
            "ARCHIVO": PDF_NAME,
            "RUTA_ELIMINADA": str(pdf_path),
            "SHA256_HISTORICO": pdf_hash,
            "TAMANO_BYTES": pdf_size,
            "ORIGINAL_DESCARGAS_CONSERVADO": "SI",
            "RUTA_ORIGINAL_DESCARGAS": str(DOWNLOADS_PDF),
            "FUENTE_NORMATIVA_ACTIVA": TXT_NAME,
            "MOTIVO": "PDF_NO_REQUERIDO_FUENTE_TXT_UNICA",
            "EJECUTOR": "Codex",
            "ESTADO": "EJECUTADO_VALIDADO",
            "OBSERVACIONES": f"Comparación TXT/PDF previa retirada: {comparison_removed}. No se ejecutó OCR ni extractor PDF.",
        }
    ]
    audit_fields = [
        "ID_ACCION",
        "FECHA_HORA",
        "ZONA_HORARIA",
        "PROCESO",
        "SUBPROCESO",
        "DECISION_INTERNA_ID",
        "ARCHIVO",
        "RUTA_ELIMINADA",
        "SHA256_HISTORICO",
        "TAMANO_BYTES",
        "ORIGINAL_DESCARGAS_CONSERVADO",
        "RUTA_ORIGINAL_DESCARGAS",
        "FUENTE_NORMATIVA_ACTIVA",
        "MOTIVO",
        "EJECUTOR",
        "ESTADO",
        "OBSERVACIONES",
    ]
    write_tsv(ROOT / "08_auditorias" / "RETIRO_PDF_INSTRUCTIVO_20260626.tsv", audit_rows, audit_fields)

    decision_md = f"""# Decisión Interna {DECISION_ID}

## Decisión

Retirar del subproyecto Avance Curricular 2026 la copia congelada del archivo `{PDF_NAME}` y establecer `{TXT_NAME}` como fuente normativa única operativa.

## Fecha

{retiro_time} America/Santiago.

## Proceso

{PROCESS}; subprocesos Carreras Avance Curricular 2026 y Matrícula Avance Curricular 2026.

## Alcance

La decisión aplica solo al subproyecto `avance_curricular_2026`. No elimina ni modifica el archivo original ubicado en Descargas.

## Justificación

La versión TXT permite búsqueda textual, extracción estructurada, trazabilidad por sección y auditoría reproducible. No se requiere comparación TXT/PDF, OCR ni herramientas de extracción PDF.

## Archivo Retirado

- Copia retirada: `{pdf_path}`
- Hash histórico: `{pdf_hash}`
- Tamaño: `{pdf_size}` bytes

## Archivo Que Permanece

- Fuente normativa única: `{txt_path}`
- Hash antes: `{before['txt']}`
- Hash después: `{after['txt']}`

## Impacto Normativo

La decisión define el soporte documental operativo. No modifica el contenido normativo extraído desde el instructivo TXT.

## Impacto Técnico

B02 queda cerrado como `DESCARTADO_POR_DECISION_INTERNA`. No hay bloqueo por extracción PDF.

## Evidencia

Ver `avance_curricular_2026/08_auditorias/RETIRO_PDF_INSTRUCTIVO_20260626.tsv` y manifiesto actualizado.

## Responsable De La Decisión

Decisión interna indicada por la persona usuaria; ejecutada y documentada por Codex.

## Validación Final

PDF retirado solo desde la carga congelada; PDF original de Descargas conservado; TXT, precargas y Excel mantienen hash sin cambios.
"""
    (ROOT / "01_documentacion" / "DECISION_INTERNA_DI-AC-2026-001.md").write_text(decision_md, encoding="utf-8")

    readme = f"""# README Carga Congelada Rectificada - Avance Curricular SIES 2026

Proceso: {PROCESS}

Fecha de carga original: 2026-06-26T00:58:26 America/Santiago.

Fecha de retiro PDF: {retiro_time} America/Santiago.

Rama: `backup/pre-sync-fix-20260410-avance`.

Commit inicial registrado: `d9c0e32`.

## Fuentes Activas

La carga congelada activa contiene cuatro fuentes:

1. `5810_Precarga Carreras Avance Curricular 20268.csv`
2. `5809_Precarga Matrícula Avance Curricular 2026.csv`
3. `{TXT_NAME}`
4. `{PROMEDIOS_NAME}`

## Decisión Interna

El PDF `{PDF_NAME}` fue retirado por decisión interna `{DECISION_ID}`.

El TXT `{TXT_NAME}` es la fuente normativa única del proceso.

No se realizó comparación TXT/PDF en esta rectificación, no se ejecutó OCR y no se instaló software de extracción PDF.

El archivo original de Descargas no fue eliminado: `{DOWNLOADS_PDF}`.

El retiro no modifica las precargas ni el Excel.

## Hash Histórico Del PDF Retirado

- Ruta retirada: `{pdf_path}`
- SHA-256 histórico: `{pdf_hash}`
- Tamaño: {pdf_size} bytes

## Hashes De Fuentes Activas

- TXT: {after['txt']}
- Carreras: {after['carreras']}
- Matrícula: {after['matricula']}
- PROMEDIOSDEALUMNOS: {after['promedios']}

## Estado

CARGA_CONGELADA_RECTIFICADA_VALIDADA
"""
    (FREEZE / "README_CARGA_CONGELADA.md").write_text(readme, encoding="utf-8")

    update_pdf_references_in_documents(retiro_time)
    update_initial_script_policy()

    return {
        "pdf_path": str(pdf_path),
        "pdf_hash": pdf_hash,
        "pdf_size": pdf_size,
        "comparison_removed": comparison_removed,
        "txt_hash_before": before["txt"],
        "txt_hash_after": after["txt"],
        "source_hashes_after": after,
        "retiro_time": retiro_time,
        "active_sources": 4,
    }


def update_pdf_references_in_documents(retiro_time: str) -> None:
    for path in (ROOT / "01_documentacion").glob("*.md"):
        text = path.read_text(encoding="utf-8")
        original = text
        text = re.sub(
            r"El P" r"DF se conserv.*?equivalencia textual automática\.",
            f"El PDF fue retirado por decisión interna {DECISION_ID}; el TXT es la fuente normativa única y no corresponde comparar TXT/PDF.",
            text,
        )
        text = re.sub(r"pendiente comparación TXT/" r"PDF", f"descartado por decisión interna {DECISION_ID}", text)
        text = re.sub(r"instalar extractor P" r"DF", f"no corresponde instalar herramientas de extracción PDF por {DECISION_ID}", text)
        text = re.sub(r"B02: P" r"DF no extraído textualmente\.", "B02: DESCARTADO_POR_DECISION_INTERNA.", text)
        text = re.sub(
            r"No se pudo extraer texto del PDF.*?TXT/PDF\.",
            f"El PDF fue retirado por decisión interna {DECISION_ID}; no existe bloqueo por extracción PDF.",
            text,
        )
        if path.name == "01_MATRIZ_NORMATIVA.md" and DECISION_ID not in text:
            text += f"\n\nNota: el PDF fue retirado el {retiro_time} por decisión interna {DECISION_ID}; el TXT es fuente normativa única.\n"
        if text != original:
            path.write_text(text, encoding="utf-8")


def update_initial_script_policy() -> None:
    text = SCRIPT_BASE.read_text(encoding="utf-8")
    original = text
    text = re.sub(
        r'\n    \{\n        "id": "FUENTE_INSTRUCTIVO_PDF",.*?\n    \},',
        "",
        text,
        flags=re.S,
    )
    text = re.sub(
        r"\ndef compare_instructivo\(txt_path: Path, pdf_path: Path\) -> dict\[str, Any\]:.*?\n\ndef build_manifest_rows",
        f"""\n\ndef compare_instructivo(txt_path: Path, pdf_path: Path) -> dict[str, Any]:\n    return {{\n        \"decision_interna\": \"{DECISION_ID}\",\n        \"estado\": \"PDF_RETIRADO_NO_COMPARAR\",\n        \"dictamen\": \"El TXT es fuente normativa única; no corresponde comparar TXT/PDF.\",\n    }}\n\n\ndef build_manifest_rows""",
        text,
        flags=re.S,
    )
    text = re.sub(
        r"- El P" r"DF se conserv.*?equivalencia textual automática\.",
        f"- El PDF fue retirado por decisión interna {DECISION_ID}; el TXT es fuente normativa única.",
        text,
    )
    text = text.replace(
        "No se pudo extraer texto del PDF con herramientas disponibles; no se asume equivalencia TXT/PDF.",
        f"El PDF fue retirado por decisión interna {DECISION_ID}; no existe bloqueo por extracción PDF.",
    )
    text = text.replace(
        "- Comparación TXT/PDF: registrada en `evidencia/COMPARACION_INSTRUCTIVO_TXT_" + "PDF.json`.",
        f"- Decisión {DECISION_ID}: no corresponde comparación TXT/PDF; TXT fuente normativa única.",
    )
    text = re.sub(
        r"\n    txt_copy = copies.get\(\"FUENTE_INSTRUCTIVO_TXT\"\).*?comparison = \{\"comparacion_contenido_textual\": \{\"estado\": \"NO_APLICA\"\}\}",
        f'\n    comparison = {{"estado": "PDF_RETIRADO_NO_COMPARAR", "decision_interna": "{DECISION_ID}"}}',
        text,
        flags=re.S,
    )
    if text != original:
        SCRIPT_BASE.write_text(text, encoding="utf-8")


def field_status(subprocess_: str, field: str) -> str:
    if subprocess_ == "Carreras Avance Curricular 2026":
        if field in {
            "CODIGO_UNICO",
            "NOMBRE_SEDE",
            "NOMBRE_CARRERA",
            "JORNADA",
            "VERSION",
            "DURACION_ESTUDIOS",
            "DURACION_TITULACION",
            "DURACION_TOTAL",
            "NIVEL_CARRERA",
        }:
            return "DEFINIDO_OFICIALMENTE"
        if field in {"PLAN_ESTUDIOS", "VIGENCIA"}:
            return "DEFINIDO_PARCIALMENTE"
        return "PENDIENTE_FUENTE_INSTITUCIONAL"
    if field in {
        "TIPO_DOCUMENTO",
        "NUM_DOCUMENTO",
        "DV",
        "PRIMER_APELLIDO",
        "SEGUNDO_APELLIDO",
        "NOMBRES",
        "SEXO",
        "FECHA_NACIMIENTO",
        "CODIGO_UNICO",
        "ANIO_INGRESO_CARRERA_ACTUAL",
        "SEM_INGRESO_CARRERA_ACTUAL",
        "ANIO_INGRESO_CARRERA_ORIGEN",
        "SEM_INGRESO_CARRERA_ORIGEN",
    }:
        return "DEFINIDO_OFICIALMENTE"
    if field in {"PLAN_ESTUDIOS", "VIGENCIA"}:
        return "DEFINIDO_PARCIALMENTE"
    return "PENDIENTE_FUENTE_INSTITUCIONAL"


def source_expected(subprocess_: str, field: str) -> tuple[str, str, str, str, str]:
    if subprocess_.startswith("Carreras"):
        if field == "PLAN_ESTUDIOS":
            return ("catálogo de planes de estudio vigentes", "PROMEDIOSDEALUMNOS_7804.xlsx", "Hoja1/matriz", "PLAN_DE_ESTUDIO/CODIGO_UNICO", "CODIGO_UNICO + evidencia institucional")
        if field in {"TIPO_UNIDAD_MEDIDA", "OTRA_UNIDAD_MEDIDA", "TOTAL_UNIDADES_MEDIDA"} or field.startswith("UNIDADES_"):
            return ("malla curricular o plan oficial por carrera", "NO_IDENTIFICADO", "NO_APLICA", "NO_IDENTIFICADO", "CODIGO_UNICO + PLAN_ESTUDIOS")
        if field == "VIGENCIA":
            return ("criterio institucional de pertenencia a la carga", "NO_IDENTIFICADO", "NO_APLICA", "NO_IDENTIFICADO", "CODIGO_UNICO + PLAN_ESTUDIOS")
        return ("precarga oficial congelada", "precarga carreras", "NO_APLICA", field, "CODIGO_UNICO")
    if field == "PLAN_ESTUDIOS":
        return ("plan real de cada estudiante", "PROMEDIOSDEALUMNOS_7804.xlsx", "Hoja1", "PLAN_DE_ESTUDIO", "CODCLI/RUT + CODCARR/CODIGO_UNICO")
    if field in {"CURSO_1ER_SEM", "CURSO_2DO_SEM"}:
        return ("inscripción o actividad académica 2025 por semestre", "PROMEDIOSDEALUMNOS_7804.xlsx", "Hoja1", "ANO + PERIODO + ESTADO", "CODCLI/RUT + CODCARR")
    if field in {"UNIDADES_CURSADAS", "UNIDADES_APROBADAS", "UNID_CURSADAS_TOTAL", "UNID_APROBADAS_TOTAL"}:
        return ("historial académico con unidades del plan", "PROMEDIOSDEALUMNOS_7804.xlsx", "Hoja1", "ASIGNATURA/NOTA_FINAL/ESTADO/CONVALIDADO", "CODCLI/RUT + CODCARR + asignatura")
    if field == "VIGENCIA":
        return ("criterio institucional de pertenencia a la carga", "NO_IDENTIFICADO", "NO_APLICA", "NO_IDENTIFICADO", "documento + CODIGO_UNICO")
    return ("precarga oficial congelada", "precarga matrícula", "NO_APLICA", field, "documento + CODIGO_UNICO")


def validation_text(subprocess_: str, field: str) -> tuple[str, str, str]:
    if field == "PLAN_ESTUDIOS":
        return (
            "Debe ser obligatorio y corresponder a plan vigente/correlativo para el CODIGO_UNICO.",
            "No asignar por defecto sin evidencia; validar existencia en Carreras para Matrícula.",
            "BLOCKER",
        )
    if field in {"TIPO_UNIDAD_MEDIDA", "OTRA_UNIDAD_MEDIDA", "TOTAL_UNIDADES_MEDIDA"} or field.startswith("UNIDADES_"):
        return (
            "Debe respetar dominio, enteros, duración y suma contra total del plan.",
            "Bloquear si falta malla o unidad institucional; no ajustar automáticamente.",
            "BLOCKER",
        )
    if field in {"UNIDADES_CURSADAS", "UNIDADES_APROBADAS", "UNID_CURSADAS_TOTAL", "UNID_APROBADAS_TOTAL"}:
        return (
            "Debe separar avance anual y acumulado; validar aprobadas <= cursadas.",
            "No incluir convalidaciones en anual; distinguir acumulado con respaldo.",
            "BLOCKER",
        )
    if field in {"CURSO_1ER_SEM", "CURSO_2DO_SEM"}:
        return (
            "Dominio SI/NO; presencia efectiva en semestre 2025.",
            "No inferir NO por ausencia de nota ni SI solo por matrícula.",
            "ERROR",
        )
    if field == "VIGENCIA":
        return (
            "Dominio 0/1; mantiene o elimina registro de esta carga.",
            "No trasladar reglas de otros procesos; VIGENCIA=0 requiere evidencia.",
            "ERROR",
        )
    if field in PERSONAL_FIELD_NAMES:
        return (
            "Campo precargado no modificable con dato personal.",
            "Preservar exactamente y mantener fuera de Git si aparece en datos reales.",
            "BLOCKER",
        )
    return ("Campo precargado no modificable.", "Comparar contra precarga congelada.", "ERROR")


def governance_rows(subprocess_: str, fields: list[str], modifiable: set[str], profile_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    profile = json.loads(profile_path.read_text(encoding="utf-8"))
    type_observed = profile.get("tipos_observados_por_columna", {})
    empty_counts = profile.get("vacios_por_columna", {})
    rows: list[dict[str, Any]] = []
    index_rows: list[dict[str, Any]] = []
    folder = "carreras" if subprocess_.startswith("Carreras") else "matricula"
    section = "Anexo I" if folder == "carreras" else "Anexo II"
    page = "8-10" if folder == "carreras" else "11-14"
    for pos, field in enumerate(fields, start=1):
        info = FIELD_DEFINITIONS[field]
        expected, archivo, hoja, origen, llave = source_expected(subprocess_, field)
        status = field_status(subprocess_, field)
        semval, block_rule, severity = validation_text(subprocess_, field)
        pending = "NO_APLICA" if status == "DEFINIDO_OFICIALMENTE" else "información no entregada por la fuente oficial; requiere fuente institucional."
        apt = "SI" if status == "DEFINIDO_OFICIALMENTE" else ("PARCIAL" if status == "DEFINIDO_PARCIALMENTE" else "NO")
        precargado = "NO" if field in modifiable and field not in {"PLAN_ESTUDIOS", "VIGENCIA"} else "SI"
        if field in {
            "TIPO_UNIDAD_MEDIDA",
            "OTRA_UNIDAD_MEDIDA",
            "TOTAL_UNIDADES_MEDIDA",
            "UNIDADES_1ER_ANIO",
            "UNIDADES_2DO_ANIO",
            "UNIDADES_3ER_ANIO",
            "UNIDADES_4TO_ANIO",
            "UNIDADES_5TO_ANIO",
            "UNIDADES_6TO_ANIO",
            "UNIDADES_7MO_ANIO",
            "CURSO_1ER_SEM",
            "CURSO_2DO_SEM",
            "UNIDADES_CURSADAS",
            "UNIDADES_APROBADAS",
            "UNID_CURSADAS_TOTAL",
            "UNID_APROBADAS_TOTAL",
        }:
            precargado = "NO"
        gov = {
            "PROCESO": PROCESS,
            "SUBPROCESO": subprocess_,
            "ANIO_PROCESO": YEAR_PROCESS,
            "ANIO_DATOS": YEAR_DATA,
            "PERIODO": PERIOD,
            "POSICION": pos,
            "NOMBRE_CAMPO": field,
            "NOMBRE_DESCRIPTIVO": info["desc"],
            "DEFINICION_OFICIAL": info["def"],
            "FUENTE_NORMATIVA": TXT_NAME,
            "SECCION": section,
            "PAGINA": page,
            "NIVEL_RESPALDO": "REGLA_OFICIAL_TXT" if status != "PENDIENTE_FUENTE_INSTITUCIONAL" else "REGLA_OFICIAL_TXT_MAS_FUENTE_INSTITUCIONAL_PENDIENTE",
            "PRECARGADO": precargado,
            "MODIFICABLE": "SI" if field in modifiable else "NO",
            "OBLIGATORIO": "SI",
            "TIPO_DATO_OFICIAL": info["type"],
            "TIPO_DATO_OBSERVADO": type_observed.get(field, "NO_OBSERVADO_EN_PRECARGA"),
            "LONGITUD_O_FORMATO": "NO_DETERMINADO_POR_INSTRUCTIVO",
            "DOMINIO": info["domain"],
            "CODIGOS_VALIDOS": info["domain"] if any(token in info["domain"] for token in [";", "0 ", "1 ", "SI"]) else "NO_APLICA",
            "NULO_PERMITIDO": "NO_DETERMINADO",
            "VACIO_PERMITIDO": "NO" if field in modifiable else "SEGUN_PRECARGA",
            "CERO_PERMITIDO": "SOLO_SI_VALOR_REAL_Y_DOMINIO_LO_PERMITE",
            "NO_APLICA_PERMITIDO": "SOLO_SI_INSTRUCTIVO_O_VALIDACION_LO_PERMITE",
            "FUENTE_INSTITUCIONAL_ESPERADA": expected,
            "ARCHIVO_FUENTE_CANDIDATO": archivo,
            "HOJA_FUENTE_CANDIDATA": hoja,
            "CAMPO_ORIGEN_CANDIDATO": origen,
            "LLAVE_DE_CRUCE_CANDIDATA": llave,
            "TRANSFORMACION_PERMITIDA": "Preservar texto; completar campos modificables solo con fuente trazable.",
            "TRANSFORMACION_PROHIBIDA": "Inferir por ausencia, usar ceros por defecto, copiar reglas de otros procesos, corregir precarga no modificable.",
            "NORMALIZACION_PERMITIDA": "Solo normalización técnica reversible para cruces documentados.",
            "NORMALIZACION_PROHIBIDA": "Cambiar códigos, fechas, nombres, planes o unidades sin respaldo oficial/institucional.",
            "DEPENDENCIAS": "Carreras validada" if subprocess_.startswith("Matrícula") and field in {"CODIGO_UNICO", "PLAN_ESTUDIOS"} else "NO_APLICA",
            "VALIDACION_ESTRUCTURAL": "Presencia de columna, posición y orden oficial.",
            "VALIDACION_SEMANTICA": semval,
            "VALIDACION_CRUZADA": "Ver RELACIONES_COLUMNAS.tsv" if field in {"CODIGO_UNICO", "PLAN_ESTUDIOS", "TIPO_UNIDAD_MEDIDA", "TOTAL_UNIDADES_MEDIDA"} or "UNID" in field else "NO_APLICA",
            "REGLA_DE_BLOQUEO": block_rule,
            "SEVERIDAD": severity,
            "SALIDA_DE_AUDITORIA": "SI",
            "DATO_OBSERVADO_PRECARGA": f"Columna observada; vacíos={empty_counts.get(field, 'NO_OBSERVADO')}; sin muestras reales.",
            "IMPLEMENTACION_EXISTENTE": "NO_REUTILIZAR_LEGACY; diseño documental nuevo.",
            "DECISION_INTERNA": f"{DECISION_ID}: TXT fuente normativa única; PDF retirado.",
            "HIPOTESIS_O_PENDIENTE": pending,
            "ESTADO_GOBERNANZA": status,
            "OBSERVACIONES": "No se completan datos en esta fase.",
        }
        rows.append(gov)
        rel_file = f"avance_curricular_2026/03_gobernanza_columnas/{folder}/{pos:02d}_{field}.tsv"
        index_rows.append(
            {
                "POSICION": pos,
                "NOMBRE_CAMPO": field,
                "SUBPROCESO": subprocess_,
                "ARCHIVO_GOBERNANZA": rel_file,
                "PRECARGADO": precargado,
                "MODIFICABLE": "SI" if field in modifiable else "NO",
                "OBLIGATORIO": "SI",
                "FUENTE_PRINCIPAL": TXT_NAME,
                "NIVEL_RESPALDO": gov["NIVEL_RESPALDO"],
                "DEPENDENCIAS": gov["DEPENDENCIAS"],
                "SEVERIDAD": severity,
                "ESTADO_GOBERNANZA": status,
                "PENDIENTE_PRINCIPAL": pending,
                "APTO_PARA_IMPLEMENTACION": apt,
            }
        )
    return rows, index_rows


GOV_FIELDS = [
    "PROCESO",
    "SUBPROCESO",
    "ANIO_PROCESO",
    "ANIO_DATOS",
    "PERIODO",
    "POSICION",
    "NOMBRE_CAMPO",
    "NOMBRE_DESCRIPTIVO",
    "DEFINICION_OFICIAL",
    "FUENTE_NORMATIVA",
    "SECCION",
    "PAGINA",
    "NIVEL_RESPALDO",
    "PRECARGADO",
    "MODIFICABLE",
    "OBLIGATORIO",
    "TIPO_DATO_OFICIAL",
    "TIPO_DATO_OBSERVADO",
    "LONGITUD_O_FORMATO",
    "DOMINIO",
    "CODIGOS_VALIDOS",
    "NULO_PERMITIDO",
    "VACIO_PERMITIDO",
    "CERO_PERMITIDO",
    "NO_APLICA_PERMITIDO",
    "FUENTE_INSTITUCIONAL_ESPERADA",
    "ARCHIVO_FUENTE_CANDIDATO",
    "HOJA_FUENTE_CANDIDATA",
    "CAMPO_ORIGEN_CANDIDATO",
    "LLAVE_DE_CRUCE_CANDIDATA",
    "TRANSFORMACION_PERMITIDA",
    "TRANSFORMACION_PROHIBIDA",
    "NORMALIZACION_PERMITIDA",
    "NORMALIZACION_PROHIBIDA",
    "DEPENDENCIAS",
    "VALIDACION_ESTRUCTURAL",
    "VALIDACION_SEMANTICA",
    "VALIDACION_CRUZADA",
    "REGLA_DE_BLOQUEO",
    "SEVERIDAD",
    "SALIDA_DE_AUDITORIA",
    "DATO_OBSERVADO_PRECARGA",
    "IMPLEMENTACION_EXISTENTE",
    "DECISION_INTERNA",
    "HIPOTESIS_O_PENDIENTE",
    "ESTADO_GOBERNANZA",
    "OBSERVACIONES",
]


INDEX_FIELDS = [
    "POSICION",
    "NOMBRE_CAMPO",
    "SUBPROCESO",
    "ARCHIVO_GOBERNANZA",
    "PRECARGADO",
    "MODIFICABLE",
    "OBLIGATORIO",
    "FUENTE_PRINCIPAL",
    "NIVEL_RESPALDO",
    "DEPENDENCIAS",
    "SEVERIDAD",
    "ESTADO_GOBERNANZA",
    "PENDIENTE_PRINCIPAL",
    "APTO_PARA_IMPLEMENTACION",
]


def phase_04_governance() -> dict[str, Counter]:
    carreras_rows, carreras_index = governance_rows(
        "Carreras Avance Curricular 2026",
        CARRERAS_FIELDS,
        CARRERAS_MODIFIABLE,
        FREEZE / "perfiles" / "PERFIL_PRECARGA_CARRERAS.json",
    )
    matricula_rows, matricula_index = governance_rows(
        "Matrícula Avance Curricular 2026",
        MATRICULA_FIELDS,
        MATRICULA_MODIFIABLE,
        FREEZE / "perfiles" / "PERFIL_PRECARGA_MATRICULA.json",
    )
    for row in carreras_rows:
        path = ROOT / "03_gobernanza_columnas" / "carreras" / f"{int(row['POSICION']):02d}_{row['NOMBRE_CAMPO']}.tsv"
        write_tsv(path, [row], GOV_FIELDS)
    for row in matricula_rows:
        path = ROOT / "03_gobernanza_columnas" / "matricula" / f"{int(row['POSICION']):02d}_{row['NOMBRE_CAMPO']}.tsv"
        write_tsv(path, [row], GOV_FIELDS)
    write_tsv(ROOT / "03_gobernanza_columnas" / "_INDICE_COLUMNAS_CARRERAS.tsv", carreras_index, INDEX_FIELDS)
    write_tsv(ROOT / "03_gobernanza_columnas" / "_INDICE_COLUMNAS_MATRICULA.tsv", matricula_index, INDEX_FIELDS)
    return {
        "carreras": Counter(row["ESTADO_GOBERNANZA"] for row in carreras_rows),
        "matricula": Counter(row["ESTADO_GOBERNANZA"] for row in matricula_rows),
    }


def validation_row(vid: str, sub: str, field: str, typ: str, desc: str, expr: str, section: str, sev: str, block_control: str, block_pes: str, status: str = "DISEÑADA") -> dict[str, str]:
    return {
        "ID_VALIDACION": vid,
        "PROCESO": PROCESS,
        "SUBPROCESO": sub,
        "CAMPO": field,
        "TIPO_VALIDACION": typ,
        "DESCRIPCION": desc,
        "EXPRESION_FUNCIONAL": expr,
        "FUENTE_REGLA": TXT_NAME,
        "SECCION": section,
        "NIVEL_RESPALDO": "REGLA_OFICIAL_EXPLICITA" if "Anexo" in section else "VALIDACION_TECNICA_DERIVADA",
        "SEVERIDAD": sev,
        "BLOQUEA_CONTROL": block_control,
        "BLOQUEA_PES": block_pes,
        "SALIDA_ERROR": f"{vid}_errores.tsv",
        "ESTADO_IMPLEMENTACION": status,
        "OBSERVACIONES": "Diseño documental; no ejecutado productivamente.",
    }


VAL_FIELDS = [
    "ID_VALIDACION",
    "PROCESO",
    "SUBPROCESO",
    "CAMPO",
    "TIPO_VALIDACION",
    "DESCRIPCION",
    "EXPRESION_FUNCIONAL",
    "FUENTE_REGLA",
    "SECCION",
    "NIVEL_RESPALDO",
    "SEVERIDAD",
    "BLOQUEA_CONTROL",
    "BLOQUEA_PES",
    "SALIDA_ERROR",
    "ESTADO_IMPLEMENTACION",
    "OBSERVACIONES",
]


def phase_04b_validations() -> dict[str, Counter]:
    car = [
        validation_row("CAR-EST-001", "Carreras Avance Curricular 2026", "TODOS", "ESTRUCTURAL", "Estructura oficial y orden de columnas.", "Columnas coinciden con contrato Carreras; CODIGO_IES_NUM solo como columna institucional previa.", "Anexo I/III", "BLOCKER", "SI", "SI"),
        validation_row("CAR-MOD-001", "Carreras Avance Curricular 2026", "NO_MODIFICABLES", "MODIFICABILIDAD", "No modificar campos precargados no modificables.", "Comparar contra precarga congelada.", "Anexo I/IV", "BLOCKER", "SI", "SI"),
        validation_row("CAR-PLAN-001", "Carreras Avance Curricular 2026", "PLAN_ESTUDIOS", "CONSISTENCIA", "Plan correlativo por CODIGO_UNICO.", "Planes múltiples solo con unidad o total distintos y evidencia institucional.", "Anexo I", "BLOCKER", "SI", "SI"),
        validation_row("CAR-UM-001", "Carreras Avance Curricular 2026", "TIPO_UNIDAD_MEDIDA", "DOMINIO", "Tipo de unidad permitido.", "Valor en {1,2,3}.", "Anexo I/IV", "BLOCKER", "SI", "SI"),
        validation_row("CAR-UM-002", "Carreras Avance Curricular 2026", "OTRA_UNIDAD_MEDIDA", "UNIDAD_MEDIDA", "Otra unidad solo cuando tipo=3.", "Si tipo=3 debe especificarse; si tipo!=3 revisar vacío.", "Anexo I/IV", "ERROR", "SI", "SI"),
        validation_row("CAR-TOT-001", "Carreras Avance Curricular 2026", "TOTAL_UNIDADES_MEDIDA", "OBLIGATORIEDAD", "Total obligatorio y entero.", "No derivar desde duración sin fuente institucional.", "Anexo I", "BLOCKER", "SI", "SI"),
        validation_row("CAR-DIST-001", "Carreras Avance Curricular 2026", "UNIDADES_1ER_ANIO..UNIDADES_7MO_ANIO", "CONSISTENCIA", "Distribución anual suma contra total.", "Suma anual = TOTAL_UNIDADES_MEDIDA; diferencias quedan bloqueadas/pendientes.", "Anexo IV", "BLOCKER", "SI", "SI"),
        validation_row("CAR-DUR-001", "Carreras Avance Curricular 2026", "UNIDADES_1ER_ANIO..UNIDADES_7MO_ANIO", "TEMPORALIDAD", "No informar años posteriores a duración formal.", "Usar DURACION_ESTUDIOS para determinar años esperados.", "Anexo IV", "ERROR", "SI", "SI"),
        validation_row("CAR-VIG-001", "Carreras Avance Curricular 2026", "VIGENCIA", "DOMINIO", "Vigencia solo 0 o 1.", "0 elimina; 1 mantiene en esta carga.", "Anexo I/IV", "ERROR", "SI", "SI"),
        validation_row("CAR-PRI-001", "Carreras Avance Curricular 2026", "TODOS", "PRIVACIDAD", "No preparar datos personales para Git público.", "Mantener fuentes locales y versionar solo artefactos saneados.", "Control interno", "BLOCKER", "SI", "SI"),
    ]
    mat = [
        validation_row("MAT-EST-001", "Matrícula Avance Curricular 2026", "TODOS", "ESTRUCTURAL", "Estructura oficial y orden de columnas.", "Columnas coinciden con contrato Matrícula; CODIGO_IES_NUM solo como columna institucional previa.", "Anexo II/III", "BLOCKER", "SI", "SI"),
        validation_row("MAT-MOD-001", "Matrícula Avance Curricular 2026", "NO_MODIFICABLES", "MODIFICABILIDAD", "No modificar atributos personales ni campos precargados.", "Comparar contra precarga congelada.", "Anexo II/V", "BLOCKER", "SI", "SI"),
        validation_row("MAT-DOC-001", "Matrícula Avance Curricular 2026", "TIPO_DOCUMENTO/DV", "DOMINIO", "Documento y DV según dominio.", "TIPO_DOCUMENTO R/P; DV según reglas del instructivo.", "Anexo II/V", "ERROR", "SI", "SI"),
        validation_row("MAT-SEX-001", "Matrícula Avance Curricular 2026", "SEXO", "DOMINIO", "Sexo H/M/X.", "No corregir directamente cambios a no binario sin procedimiento oficial.", "Anexo II/V", "ERROR", "SI", "SI"),
        validation_row("MAT-PLAN-001", "Matrícula Avance Curricular 2026", "PLAN_ESTUDIOS", "CRUZADA", "Plan debe existir en Carreras.", "CODIGO_UNICO + PLAN_ESTUDIOS debe existir en catálogo validado.", "Anexo II/V", "BLOCKER", "SI", "SI"),
        validation_row("MAT-CURSO-001", "Matrícula Avance Curricular 2026", "CURSO_1ER_SEM", "DOMINIO", "Valores SI/NO.", "No inferir por ausencia de nota.", "Anexo II/V", "ERROR", "SI", "SI"),
        validation_row("MAT-CURSO-002", "Matrícula Avance Curricular 2026", "CURSO_2DO_SEM", "DOMINIO", "Valores SI/NO.", "No inferir por ausencia de nota.", "Anexo II/V", "ERROR", "SI", "SI"),
        validation_row("MAT-ANUAL-001", "Matrícula Avance Curricular 2026", "UNIDADES_APROBADAS", "CONSISTENCIA", "Aprobadas anuales no superan cursadas anuales.", "UNIDADES_APROBADAS <= UNIDADES_CURSADAS.", "Anexo V", "ERROR", "SI", "SI"),
        validation_row("MAT-ACUM-001", "Matrícula Avance Curricular 2026", "UNID_APROBADAS_TOTAL", "CONSISTENCIA", "Aprobadas total no superan cursadas total.", "UNID_APROBADAS_TOTAL <= UNID_CURSADAS_TOTAL.", "Anexo V", "ERROR", "SI", "SI"),
        validation_row("MAT-ACUM-002", "Matrícula Avance Curricular 2026", "UNID_APROBADAS_TOTAL", "UNIDAD_MEDIDA", "Comparación contra total del plan.", "Aplicar control de referencia y tolerancia 25% como alerta/bloqueo diseñado.", "3.2/Anexo V", "WARNING", "NO", "SI"),
        validation_row("MAT-CONV-001", "Matrícula Avance Curricular 2026", "UNIDADES_CURSADAS/UNIDADES_APROBADAS", "TRAZABILIDAD", "Anual no incluye convalidaciones ni reconocimiento.", "Separar anual de acumulado.", "3.2", "BLOCKER", "SI", "SI"),
        validation_row("MAT-VIG-001", "Matrícula Avance Curricular 2026", "VIGENCIA", "DOMINIO", "Vigencia solo 0 o 1 y no representa estado 2026.", "Mantener 1 para universo 2025 salvo error documentado.", "Página 7/Anexo V", "ERROR", "SI", "SI"),
        validation_row("MAT-PRI-001", "Matrícula Avance Curricular 2026", "DATOS_PERSONALES", "PRIVACIDAD", "Datos personales fuera de Git público.", "No usar muestras reales en documentación o tests.", "Control interno", "BLOCKER", "SI", "SI"),
    ]
    cross = [
        validation_row("CRZ-001", "Ambos", "CODIGO_UNICO", "CRUZADA", "Cada CODIGO_UNICO de Matrícula existe en Carreras.", "matricula.CODIGO_UNICO in carreras.CODIGO_UNICO.", "Anexo III/V", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-002", "Ambos", "CODIGO_UNICO+PLAN_ESTUDIOS", "CRUZADA", "Cada plan de Matrícula existe en Carreras.", "matricula.(CODIGO_UNICO,PLAN_ESTUDIOS) in carreras.", "Anexo II/V", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-003", "Carreras", "PLAN_ESTUDIOS", "CONSISTENCIA", "Planes correlativos por CODIGO_UNICO.", "Sin saltos ni duplicados contradictorios.", "Anexo I", "ERROR", "SI", "SI"),
        validation_row("CRZ-004", "Ambos", "TIPO_UNIDAD_MEDIDA", "UNIDAD_MEDIDA", "Matrícula usa unidad definida para el plan.", "No mezclar créditos/asignaturas.", "Página 2/3.1/3.2", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-005", "Matrícula", "UNID_APROBADAS_TOTAL", "CRUZADA", "Comparar aprobado total contra total del plan.", "Usar TOTAL_UNIDADES_MEDIDA y tolerancia 25%.", "3.2/Anexo V", "WARNING", "NO", "SI"),
        validation_row("CRZ-006", "Matrícula", "CURSO_1ER_SEM/CURSO_2DO_SEM", "CONSISTENCIA", "Presencia coherente con unidades anuales.", "Ambos NO con unidades >0 genera revisión, no corrección automática.", "Anexo V", "WARNING", "NO", "SI"),
        validation_row("CRZ-007", "Matrícula", "UNIDADES_CURSADAS", "TRAZABILIDAD", "No reemplazar desconocidos por cero.", "Distinguir cero real, vacío, no aplica, desconocido y pendiente.", "Control interno derivado", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-008", "Matrícula", "AVANCE", "TEMPORALIDAD", "Avance corresponde a carrera informada en 2025.", "No trasladar avance de cambio interno posterior.", "Página 5", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-009", "Ambos", "NO_MODIFICABLES", "MODIFICABILIDAD", "Campos precargados no modificables preservados.", "Comparar contra precarga congelada.", "Anexos I-II", "BLOCKER", "SI", "SI"),
        validation_row("CRZ-010", "Ambos", "CODIGO_IES_NUM", "ESTRUCTURAL", "Primera columna institucional solo se elimina en PES final.", "Conservar en control y congelados.", "Anexo III", "ERROR", "SI", "SI"),
    ]
    write_tsv(ROOT / "04_configuracion" / "CATALOGO_VALIDACIONES_CARRERAS.tsv", car, VAL_FIELDS)
    write_tsv(ROOT / "04_configuracion" / "CATALOGO_VALIDACIONES_MATRICULA.tsv", mat, VAL_FIELDS)
    write_tsv(ROOT / "04_configuracion" / "CATALOGO_VALIDACIONES_CRUZADAS.tsv", cross, VAL_FIELDS)
    write_relaciones()
    all_rows = car + mat + cross
    return {
        "carreras": Counter(row["SEVERIDAD"] for row in car),
        "matricula": Counter(row["SEVERIDAD"] for row in mat),
        "cruzadas": Counter(row["SEVERIDAD"] for row in cross),
        "total": Counter(row["SEVERIDAD"] for row in all_rows),
    }


def write_relaciones() -> None:
    rows = [
        ("REL-001", "Matrícula", "CODIGO_UNICO", "Carreras", "CODIGO_UNICO", "EXISTENCIA", "N:1", "Matrícula.CODIGO_UNICO debe existir en Carreras.CODIGO_UNICO", "Anexo III/V", "BLOCKER"),
        ("REL-002", "Matrícula", "CODIGO_UNICO+PLAN_ESTUDIOS", "Carreras", "CODIGO_UNICO+PLAN_ESTUDIOS", "EXISTENCIA", "N:1", "Matrícula.CODIGO_UNICO + PLAN_ESTUDIOS debe existir en Carreras", "Anexo II/V", "BLOCKER"),
        ("REL-003", "Carreras", "TIPO_UNIDAD_MEDIDA", "Matrícula", "UNIDADES_*", "UNIDAD_MEDIDA", "1:N", "Matrícula utiliza unidad de medida del plan", "Página 2", "BLOCKER"),
        ("REL-004", "Matrícula", "UNID_APROBADAS_TOTAL", "Carreras", "TOTAL_UNIDADES_MEDIDA", "REFERENCIA", "N:1", "Aprobado total se compara con total del plan", "3.2", "WARNING"),
        ("REL-005", "Carreras", "UNIDADES_1ER_ANIO..UNIDADES_7MO_ANIO", "Carreras", "TOTAL_UNIDADES_MEDIDA", "SUMA", "1:1", "Distribución anual se compara con total", "Anexo IV", "BLOCKER"),
        ("REL-006", "Matrícula", "UNIDADES_APROBADAS", "Matrícula", "UNIDADES_CURSADAS", "DESIGUALDAD", "1:1", "UNIDADES_APROBADAS <= UNIDADES_CURSADAS", "Anexo V", "ERROR"),
        ("REL-007", "Matrícula", "UNID_APROBADAS_TOTAL", "Matrícula", "UNID_CURSADAS_TOTAL", "DESIGUALDAD", "1:1", "UNID_APROBADAS_TOTAL <= UNID_CURSADAS_TOTAL", "Anexo V", "ERROR"),
        ("REL-008", "Matrícula", "CONVALIDACIONES", "Matrícula", "UNIDADES_CURSADAS/UNIDADES_APROBADAS", "EXCLUSION_ANUAL", "N:1", "Datos anuales no incluyen convalidaciones", "3.2", "BLOCKER"),
        ("REL-009", "Matrícula", "CONVALIDACIONES", "Matrícula", "UNID_CURSADAS_TOTAL/UNID_APROBADAS_TOTAL", "INCLUSION_ACUMULADO", "N:1", "Datos acumulados sí pueden incluirlas con respaldo", "3.2", "ERROR"),
        ("REL-010", "Matrícula", "CURSO_1ER_SEM+CURSO_2DO_SEM", "Matrícula", "UNIDADES_CURSADAS", "COHERENCIA", "1:1", "Ambos CURSO=NO con unidades >0 genera revisión, no corrección automática", "Anexo V", "WARNING"),
    ]
    out = [
        {
            "ID_RELACION": rid,
            "SUBPROCESO_ORIGEN": so,
            "CAMPO_ORIGEN": co,
            "SUBPROCESO_DESTINO": sd,
            "CAMPO_DESTINO": cd,
            "TIPO_RELACION": tr,
            "CARDINALIDAD_ESPERADA": card,
            "REGLA": regla,
            "FUENTE": TXT_NAME,
            "SEVERIDAD": sev,
            "ESTADO": "DISEÑADA",
            "OBSERVACIONES": "No ejecutada productivamente.",
        }
        for rid, so, co, sd, cd, tr, card, regla, fuente, sev in rows
    ]
    write_tsv(
        ROOT / "04_configuracion" / "RELACIONES_COLUMNAS.tsv",
        out,
        [
            "ID_RELACION",
            "SUBPROCESO_ORIGEN",
            "CAMPO_ORIGEN",
            "SUBPROCESO_DESTINO",
            "CAMPO_DESTINO",
            "TIPO_RELACION",
            "CARDINALIDAD_ESPERADA",
            "REGLA",
            "FUENTE",
            "SEVERIDAD",
            "ESTADO",
            "OBSERVACIONES",
        ],
    )


def normalize_col(value: str) -> str:
    value = safe_name(value)
    return value if value else "COLUMNA_SIN_NOMBRE"


def is_personal_col(col: str) -> bool:
    return normalize_col(col) in {normalize_col(x) for x in PERSONAL_FIELD_NAMES}


def masked_examples(types: Counter[str]) -> str:
    if not types:
        return "SIN_DATOS"
    labels = []
    for typ in types:
        labels.append(f"<{typ}>")
    return " | ".join(labels[:5])


def classify_value(value: Any, col: str) -> str:
    if value is None or value == "":
        return "VACIO"
    text = str(value).strip()
    ncol = normalize_col(col)
    if ncol in {normalize_col(x) for x in PERSONAL_FIELD_NAMES}:
        return "DATO_PERSONAL_ENMASCARADO"
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}.*", text) or re.fullmatch(r"\d{2}-\d{2}-\d{4}", text):
        return "FECHA"
    if re.fullmatch(r"-?\d+([.,]\d+)?", text):
        return "NUMERO"
    if text.upper() in {"SI", "NO"}:
        return "SI_NO"
    if re.fullmatch(r"[A-Z0-9_ -]{3,30}", text.upper()):
        return "CODIGO_O_TEXTO_CORTO"
    return "TEXTO"


def possible_meaning(col: str) -> str:
    n = normalize_col(col)
    mapping = {
        "CODCLI": "identificador interno de estudiante/carrera",
        "RUT": "identificador personal",
        "DIG": "dígito verificador",
        "DV": "dígito verificador",
        "CODCARR": "código interno de carrera",
        "CODCARPR": "código interno de carrera/programa",
        "CODIGO_UNICO": "código único SIES",
        "PLAN_DE_ESTUDIO": "plan de estudios candidato",
        "ANO": "año académico",
        "PERIODO": "periodo/semestre",
        "ASIGNATURA": "asignatura",
        "NOTA_FINAL": "resultado académico",
        "ESTADO": "estado de asignatura o registro",
        "CONVALIDADO": "marca potencial de convalidación",
        "DURACION_ESTUDIOS": "duración de estudios",
        "DURACION_TITULACION": "duración titulación",
        "DURACION_TOTAL": "duración total",
    }
    return mapping.get(n, "NO_DETERMINADO")


def destination_candidate(col: str) -> str:
    n = normalize_col(col)
    if n in {"CODIGO_UNICO"}:
        return "Carreras.CODIGO_UNICO | Matrícula.CODIGO_UNICO"
    if n in {"PLAN_DE_ESTUDIO", "PLAN_ESTUDIOS"}:
        return "PLAN_ESTUDIOS (requiere validación)"
    if n in {"ANO", "PERIODO", "ANOMATRICULA", "PERIODOMATRICULA"}:
        return "CURSO_1ER_SEM/CURSO_2DO_SEM o filtros temporales"
    if n in {"ASIGNATURA", "CODRAMO", "RAMOEQUIV"}:
        return "UNIDADES_CURSADAS/UNID_CURSADAS_TOTAL (candidato)"
    if n in {"NOTA_FINAL", "ESTADO", "DESCRIPCION_ESTADO"}:
        return "UNIDADES_APROBADAS/UNID_APROBADAS_TOTAL (candidato)"
    if n in {"CONVALIDADO"}:
        return "separación anual/acumulado por convalidaciones"
    if n in {"DURACION_ESTUDIOS", "DURACION_TITULACION", "DURACION_TOTAL", "NIVEL_CARRERA", "VIGENCIA"}:
        return f"Carreras.{n}"
    return "NO_DETERMINADO"


def phase_04d_promedios() -> dict[str, Any]:
    path = ORIG / PROMEDIOS_NAME
    wb = load_workbook(path, read_only=True, data_only=False)
    sheet_rows: list[dict[str, Any]] = []
    dictionary_rows: list[dict[str, Any]] = []
    coverage_support: dict[str, dict[str, Any]] = {}
    for sidx, ws in enumerate(wb.worksheets, start=1):
        headers = [str(v) if v is not None else "" for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        nonempty_counts = [0] * len(headers)
        null_counts = [0] * len(headers)
        distinct_sets: list[set[str]] = [set() for _ in headers]
        type_counts: list[Counter[str]] = [Counter() for _ in headers]
        has_2025 = False
        max_row = ws.max_row or 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            for i, value in enumerate(row[: len(headers)]):
                if value is None or value == "":
                    null_counts[i] += 1
                else:
                    nonempty_counts[i] += 1
                    if len(distinct_sets[i]) < 10000:
                        distinct_sets[i].add(str(value))
                    typ = classify_value(value, headers[i] if i < len(headers) else "")
                    if typ != "VACIO":
                        type_counts[i][typ] += 1
                    if str(value).strip() == "2025":
                        has_2025 = True
        nheaders = [normalize_col(h) for h in headers]
        header_set = set(nheaders)
        gran = "NO_DETERMINADO"
        if {"CODCLI", "ASIGNATURA"} & header_set and {"ANO", "PERIODO"} <= header_set:
            gran = "PERSONA_ASIGNATURA"
        elif "CODCLI" in header_set and ("CODCARPR" in header_set or "CODCARR" in header_set):
            gran = "PERSONA_CARRERA"
        elif "CODIGO_UNICO" in header_set and "NOMBRE_CARRERA" in header_set:
            gran = "CARRERA_PLAN"
        elif "RUT" in header_set and "CODCLI" in header_set:
            gran = "PERSONA"
        columns_key = [h for h in headers if normalize_col(h) in {"CODCLI", "RUT", "DIG", "DV", "CODCARR", "CODCARPR", "CODIGO_UNICO", "PLAN_DE_ESTUDIO", "ANO", "PERIODO"}]
        sheet_rows.append(
            {
                "INDICE_HOJA": sidx,
                "NOMBRE_HOJA": ws.title,
                "FILAS": max_row,
                "COLUMNAS": ws.max_column,
                "FILA_ENCABEZADO": 1,
                "COLUMNAS_CLAVE": columns_key,
                "CONTIENE_CODCLI": "SI" if "CODCLI" in header_set else "NO",
                "CONTIENE_RUT": "SI" if "RUT" in header_set else "NO",
                "CONTIENE_CODIGO_CARRERA": "SI" if {"CODCARR", "CODCARPR", "CODIGO_UNICO"} & header_set else "NO",
                "CONTIENE_PLAN_ESTUDIOS": "SI" if {"PLAN_DE_ESTUDIO", "PLAN_ESTUDIOS"} & header_set else "NO",
                "CONTIENE_PERIODO": "SI" if {"PERIODO", "PERIODOMATRICULA"} & header_set else "NO",
                "CONTIENE_ANIO": "SI" if {"ANO", "ANOMATRICULA", "ANOINGRESO"} & header_set else "NO",
                "CONTIENE_SEMESTRE": "SI" if {"PERIODO", "PERIODOMATRICULA", "PERIODOINGRESO"} & header_set else "NO",
                "CONTIENE_ASIGNATURA": "SI" if {"ASIGNATURA", "CODRAMO", "RAMOEQUIV"} & header_set else "NO",
                "CONTIENE_CREDITOS": "SI" if any("CRED" in h for h in header_set) else "NO",
                "CONTIENE_NOTA": "SI" if any("NOTA" in h for h in header_set) else "NO",
                "CONTIENE_ESTADO_ASIGNATURA": "SI" if {"ESTADO", "DESCRIPCION_ESTADO"} & header_set else "NO",
                "CONTIENE_CONVALIDACION": "SI" if any("CONVALID" in h for h in header_set) else "NO",
                "CONTIENE_HOMOLOGACION": "SI" if any("HOMOLOG" in h for h in header_set) else "NO",
                "CONTIENE_RAP": "SI" if any(h in {"RAP", "RECONOCIMIENTO_APRENDIZAJES_PREVIOS"} for h in header_set) else "NO",
                "CONTIENE_INSCRIPCION": "SI" if {"MATRICULA", "FECHAMATRICULA"} & header_set else "NO",
                "CONTIENE_DATOS_2025": "SI" if has_2025 else "NO",
                "NIVEL_GRANULARIDAD": gran,
                "RELEVANCIA_CARRERAS": "POTENCIAL" if gran == "CARRERA_PLAN" or "PLAN_DE_ESTUDIO" in header_set else "BAJA",
                "RELEVANCIA_MATRICULA": "POTENCIAL" if gran.startswith("PERSONA") else "BAJA",
                "CONTIENE_DATOS_PERSONALES": "SI" if any(is_personal_col(h) for h in headers) else "NO",
                "ESTADO": "PERFILADO_SIN_USO_AUTOMATICO",
                "OBSERVACIONES": "Perfil sin muestras reales; requiere linaje y validación institucional.",
            }
        )
        for pos, col in enumerate(headers, start=1):
            norm = normalize_col(col or f"COLUMNA_{pos}")
            dictionary_rows.append(
                {
                    "HOJA": ws.title,
                    "POSICION": pos,
                    "COLUMNA_ORIGINAL": col or f"COLUMNA_SIN_NOMBRE_{pos}",
                    "COLUMNA_NORMALIZADA": norm,
                    "TIPO_OBSERVADO": sorted(type_counts[pos - 1]) or ["VACIO"],
                    "EJEMPLOS_ENMASCARADOS": masked_examples(type_counts[pos - 1]),
                    "NULOS": null_counts[pos - 1],
                    "NO_NULOS": nonempty_counts[pos - 1],
                    "VALORES_DISTINTOS": len(distinct_sets[pos - 1]),
                    "POSIBLE_SIGNIFICADO": possible_meaning(col),
                    "CAMPO_DESTINO_CANDIDATO": destination_candidate(col),
                    "NIVEL_RESPALDO": "DATO_OBSERVADO_FUENTE_INSTITUCIONAL_CANDIDATA",
                    "TRANSFORMACION_CANDIDATA": "Solo mapeo candidato; no apto automático sin validación.",
                    "RIESGO": "ALTO_DATOS_PERSONALES" if is_personal_col(col) else "MEDIO",
                    "ESTADO_MAPEO": "CANDIDATO" if destination_candidate(col) != "NO_DETERMINADO" else "NO_MAPEADO",
                    "OBSERVACIONES": "No contiene valores reales en este diccionario.",
                }
            )
            coverage_support[f"{ws.title}.{norm}"] = {
                "sheet": ws.title,
                "col": col,
                "nonempty": nonempty_counts[pos - 1],
                "total": max(max_row - 1, 0),
                "period_2025": has_2025,
            }
    write_tsv(ROOT / "07_control" / "PERFIL_HOJAS_PROMEDIOSDEALUMNOS.tsv", sheet_rows, [
        "INDICE_HOJA",
        "NOMBRE_HOJA",
        "FILAS",
        "COLUMNAS",
        "FILA_ENCABEZADO",
        "COLUMNAS_CLAVE",
        "CONTIENE_CODCLI",
        "CONTIENE_RUT",
        "CONTIENE_CODIGO_CARRERA",
        "CONTIENE_PLAN_ESTUDIOS",
        "CONTIENE_PERIODO",
        "CONTIENE_ANIO",
        "CONTIENE_SEMESTRE",
        "CONTIENE_ASIGNATURA",
        "CONTIENE_CREDITOS",
        "CONTIENE_NOTA",
        "CONTIENE_ESTADO_ASIGNATURA",
        "CONTIENE_CONVALIDACION",
        "CONTIENE_HOMOLOGACION",
        "CONTIENE_RAP",
        "CONTIENE_INSCRIPCION",
        "CONTIENE_DATOS_2025",
        "NIVEL_GRANULARIDAD",
        "RELEVANCIA_CARRERAS",
        "RELEVANCIA_MATRICULA",
        "CONTIENE_DATOS_PERSONALES",
        "ESTADO",
        "OBSERVACIONES",
    ])
    write_tsv(ROOT / "07_control" / "DICCIONARIO_PROMEDIOSDEALUMNOS.tsv", dictionary_rows, [
        "HOJA",
        "POSICION",
        "COLUMNA_ORIGINAL",
        "COLUMNA_NORMALIZADA",
        "TIPO_OBSERVADO",
        "EJEMPLOS_ENMASCARADOS",
        "NULOS",
        "NO_NULOS",
        "VALORES_DISTINTOS",
        "POSIBLE_SIGNIFICADO",
        "CAMPO_DESTINO_CANDIDATO",
        "NIVEL_RESPALDO",
        "TRANSFORMACION_CANDIDATA",
        "RIESGO",
        "ESTADO_MAPEO",
        "OBSERVACIONES",
    ])
    coverage_rows = build_coverage_rows(coverage_support)
    write_tsv(ROOT / "07_control" / "COBERTURA_CAMPOS_DESDE_PROMEDIOS.tsv", coverage_rows, [
        "SUBPROCESO",
        "CAMPO_DESTINO",
        "HOJA_CANDIDATA",
        "COLUMNAS_CANDIDATAS",
        "LLAVE_CANDIDATA",
        "COBERTURA_FILAS",
        "COBERTURA_PORCENTAJE",
        "CALIDAD",
        "CONFLICTOS",
        "PERIODO_DETECTADO",
        "UNIDAD_DE_MEDIDA_DETECTADA",
        "NIVEL_RESPALDO",
        "USO_PROPUESTO",
        "APTO_AUTOMATICO",
        "REQUIERE_VALIDACION",
        "ESTADO",
        "OBSERVACIONES",
    ])
    return {
        "sheets": len(sheet_rows),
        "dict_rows": len(dictionary_rows),
        "coverage_rows": len(coverage_rows),
        "sheet_granularity": {row["NOMBRE_HOJA"]: row["NIVEL_GRANULARIDAD"] for row in sheet_rows},
    }


def coverage_for(support: dict[str, Any], sheet: str, col_norm: str) -> tuple[int, str]:
    key = f"{sheet}.{col_norm}"
    item = support.get(key)
    if not item:
        return 0, "0.00"
    total = item["total"]
    pct = (item["nonempty"] / total * 100) if total else 0
    return item["nonempty"], f"{pct:.2f}"


def build_coverage_rows(support: dict[str, Any]) -> list[dict[str, Any]]:
    specs = [
        ("Carreras Avance Curricular 2026", "PLAN_ESTUDIOS", "Hoja1", "PLAN_DE_ESTUDIO", "CODCARR/PLAN_DE_ESTUDIO", "PARCIAL"),
        ("Carreras Avance Curricular 2026", "TIPO_UNIDAD_MEDIDA", "NO_IDENTIFICADA", "NO_IDENTIFICADA", "CODIGO_UNICO+PLAN_ESTUDIOS", "NO"),
        ("Carreras Avance Curricular 2026", "TOTAL_UNIDADES_MEDIDA", "NO_IDENTIFICADA", "NO_IDENTIFICADA", "CODIGO_UNICO+PLAN_ESTUDIOS", "NO"),
        ("Carreras Avance Curricular 2026", "UNIDADES_1ER_ANIO..UNIDADES_7MO_ANIO", "NO_IDENTIFICADA", "NO_IDENTIFICADA", "CODIGO_UNICO+PLAN_ESTUDIOS", "NO"),
        ("Matrícula Avance Curricular 2026", "PLAN_ESTUDIOS", "Hoja1", "PLAN_DE_ESTUDIO", "CODCLI/RUT+CODCARR", "PARCIAL"),
        ("Matrícula Avance Curricular 2026", "CURSO_1ER_SEM/CURSO_2DO_SEM", "Hoja1", "ANO+PERIODO+ESTADO", "CODCLI/RUT+CODCARR", "PARCIAL"),
        ("Matrícula Avance Curricular 2026", "UNIDADES_CURSADAS", "Hoja1", "ASIGNATURA+ANO+PERIODO", "CODCLI/RUT+CODCARR+ASIGNATURA", "PARCIAL"),
        ("Matrícula Avance Curricular 2026", "UNIDADES_APROBADAS", "Hoja1", "NOTA_FINAL+ESTADO", "CODCLI/RUT+CODCARR+ASIGNATURA", "PARCIAL"),
        ("Matrícula Avance Curricular 2026", "UNID_CURSADAS_TOTAL/UNID_APROBADAS_TOTAL", "Hoja1", "ASIGNATURA+NOTA_FINAL+ESTADO+CONVALIDADO", "CODCLI/RUT+CODCARR+ASIGNATURA", "PARCIAL"),
        ("Matrícula Avance Curricular 2026", "VIGENCIA", "NO_IDENTIFICADA", "NO_IDENTIFICADA", "documento+CODIGO_UNICO", "NO"),
    ]
    rows = []
    for sub, field, sheet, cols, key, apt in specs:
        coverage = 0
        pct = "0.00"
        if sheet != "NO_IDENTIFICADA":
            first = normalize_col(cols.split("+")[0])
            coverage, pct = coverage_for(support, sheet, first)
        rows.append(
            {
                "SUBPROCESO": sub,
                "CAMPO_DESTINO": field,
                "HOJA_CANDIDATA": sheet,
                "COLUMNAS_CANDIDATAS": cols,
                "LLAVE_CANDIDATA": key,
                "COBERTURA_FILAS": coverage,
                "COBERTURA_PORCENTAJE": pct,
                "CALIDAD": "PARCIAL_REQUIERE_LINAJE" if apt == "PARCIAL" else "NO_DISPONIBLE",
                "CONFLICTOS": "Requiere equivalencia CODCARR/CODIGO_UNICO y unidad de medida del plan." if apt == "PARCIAL" else "Fuente no identificada.",
                "PERIODO_DETECTADO": "2025 detectado en hoja candidata" if sheet == "Hoja1" else "NO_DETERMINADO",
                "UNIDAD_DE_MEDIDA_DETECTADA": "NO",
                "NIVEL_RESPALDO": "DATO_OBSERVADO_CANDIDATO",
                "USO_PROPUESTO": "Perfil y diseño; no uso automático.",
                "APTO_AUTOMATICO": apt,
                "REQUIERE_VALIDACION": "SI",
                "ESTADO": "CANDIDATO" if apt == "PARCIAL" else "FALTANTE",
                "OBSERVACIONES": "No completa datos; requiere fuentes institucionales definitivas.",
            }
        )
    return rows


def phase_04e_sources() -> dict[str, int]:
    required = []
    for field in [
        "PLAN_ESTUDIOS",
        "TIPO_UNIDAD_MEDIDA",
        "OTRA_UNIDAD_MEDIDA",
        "TOTAL_UNIDADES_MEDIDA",
        "UNIDADES_1ER_ANIO",
        "UNIDADES_2DO_ANIO",
        "UNIDADES_3ER_ANIO",
        "UNIDADES_4TO_ANIO",
        "UNIDADES_5TO_ANIO",
        "UNIDADES_6TO_ANIO",
        "UNIDADES_7MO_ANIO",
        "VIGENCIA",
    ]:
        required.append(("Carreras Avance Curricular 2026", field))
    for field in [
        "PLAN_ESTUDIOS",
        "CURSO_1ER_SEM",
        "CURSO_2DO_SEM",
        "UNIDADES_CURSADAS",
        "UNIDADES_APROBADAS",
        "UNID_CURSADAS_TOTAL",
        "UNID_APROBADAS_TOTAL",
        "VIGENCIA",
    ]:
        required.append(("Matrícula Avance Curricular 2026", field))

    matrix = []
    pending = []
    for idx, (sub, field) in enumerate(required, start=1):
        expected, archivo, hoja, origen, llave = source_expected(sub, field)
        partial = archivo == PROMEDIOS_NAME or "PROMEDIOS" in archivo
        missing = archivo == "NO_IDENTIFICADO"
        estado = "FUENTE_PARCIAL_CANDIDATA" if partial else ("FUENTE_FALTANTE" if missing else "FUENTE_PRECARGA_NO_COMPLETA_DESTINO")
        blocks = "SI" if missing or field not in {"VIGENCIA"} else "PARCIAL"
        matrix.append(
            {
                "ID": f"FC-{idx:03d}",
                "SUBPROCESO": sub,
                "CAMPO_DESTINO": field,
                "FUENTE_REQUERIDA": expected,
                "ARCHIVO_DISPONIBLE": archivo,
                "HOJA": hoja,
                "COLUMNA_ORIGEN": origen,
                "LLAVE": llave,
                "GRANULARIDAD": "CARRERA_PLAN" if sub.startswith("Carreras") else "PERSONA_CARRERA/PERSONA_ASIGNATURA",
                "PERIODO": "2025" if sub.startswith("Matrícula") else "planes vigentes 2025",
                "UNIDAD_MEDIDA": "NO_IDENTIFICADA" if "UNIDAD" in field or "UNIDADES" in field else "NO_APLICA",
                "COBERTURA": "PARCIAL" if partial else "NO_IDENTIFICADA",
                "CONFLICTOS": "Requiere validación institucional y llaves." if partial else "No hay archivo disponible suficiente.",
                "TRANSFORMACION": "Diseñar; no ejecutar.",
                "NIVEL_RESPALDO": "DATO_OBSERVADO_CANDIDATO" if partial else "PENDIENTE",
                "ESTADO": estado,
                "BLOQUEA_IMPLEMENTACION": blocks,
                "ACCION_REQUERIDA": "Aportar/validar fuente institucional definitiva.",
                "OBSERVACIONES": "No se inventan fuentes ni reglas.",
            }
        )
        if missing or partial:
            pending.append(
                {
                    "ID_PENDIENTE": f"PF-{idx:03d}",
                    "SUBPROCESO": sub,
                    "CAMPO_AFECTADO": field,
                    "FUENTE_FALTANTE": expected,
                    "POR_QUE_SE_REQUIERE": "El instructivo exige completar/validar este campo con respaldo institucional.",
                    "IMPACTO": "Bloquea procesamiento productivo o exige validación manual.",
                    "SEVERIDAD": "BLOCKER" if field not in {"VIGENCIA"} else "ERROR",
                    "BLOQUEA_CARRERAS": "SI" if sub.startswith("Carreras") else "NO",
                    "BLOQUEA_MATRICULA": "SI" if sub.startswith("Matrícula") or field in {"PLAN_ESTUDIOS", "TIPO_UNIDAD_MEDIDA", "TOTAL_UNIDADES_MEDIDA"} else "NO",
                    "ACCION_SOLICITADA": "Confirmar fuente, llave, período, unidad y regla de uso.",
                    "ESTADO": "ABIERTO",
                    "OBSERVACIONES": "PROMEDIOSDEALUMNOS puede ser parcial pero no suficiente automáticamente." if partial else "No identificada en archivos disponibles.",
                }
            )
    write_tsv(ROOT / "07_control" / "MATRIZ_FUENTES_CAMPOS.tsv", matrix, [
        "ID",
        "SUBPROCESO",
        "CAMPO_DESTINO",
        "FUENTE_REQUERIDA",
        "ARCHIVO_DISPONIBLE",
        "HOJA",
        "COLUMNA_ORIGEN",
        "LLAVE",
        "GRANULARIDAD",
        "PERIODO",
        "UNIDAD_MEDIDA",
        "COBERTURA",
        "CONFLICTOS",
        "TRANSFORMACION",
        "NIVEL_RESPALDO",
        "ESTADO",
        "BLOQUEA_IMPLEMENTACION",
        "ACCION_REQUERIDA",
        "OBSERVACIONES",
    ])
    write_cruces_propuesta()
    write_tsv(ROOT / "09_pendientes" / "FUENTES_INSTITUCIONALES_FALTANTES.tsv", pending, [
        "ID_PENDIENTE",
        "SUBPROCESO",
        "CAMPO_AFECTADO",
        "FUENTE_FALTANTE",
        "POR_QUE_SE_REQUIERE",
        "IMPACTO",
        "SEVERIDAD",
        "BLOQUEA_CARRERAS",
        "BLOQUEA_MATRICULA",
        "ACCION_SOLICITADA",
        "ESTADO",
        "OBSERVACIONES",
    ])
    return {"matrix": len(matrix), "pending": len(pending), "partial": sum(1 for r in matrix if r["ESTADO"] == "FUENTE_PARCIAL_CANDIDATA"), "missing": sum(1 for r in matrix if r["ESTADO"] == "FUENTE_FALTANTE")}


def write_cruces_propuesta() -> None:
    rows = [
        ("CRUCE-001", "Vincular matrícula a carreras", "Precarga Matrícula", "PERSONA_CARRERA", "Precarga Carreras", "CARRERA_PLAN", "CODIGO_UNICO+PLAN_ESTUDIOS", "CODIGO_UNICO+PLAN_ESTUDIOS", "N:1", "unidad, total, vigencia plan", "MEDIO"),
        ("CRUCE-002", "Resolver plan candidato", "PROMEDIOS Hoja1", "PERSONA_ASIGNATURA", "Precarga Matrícula", "PERSONA_CARRERA", "RUT/CODCLI+CODCARR", "documento+CODIGO_UNICO", "N:1 con puente validado", "PLAN_DE_ESTUDIO", "ALTO"),
        ("CRUCE-003", "Detectar presencia 2025", "PROMEDIOS Hoja1", "PERSONA_ASIGNATURA", "Precarga Matrícula", "PERSONA_CARRERA", "RUT/CODCLI+CODCARR+ANO+PERIODO", "documento+CODIGO_UNICO", "N:1 agregada", "CURSO_1ER_SEM/CURSO_2DO_SEM", "ALTO"),
        ("CRUCE-004", "Calcular avance anual candidato", "PROMEDIOS Hoja1", "PERSONA_ASIGNATURA", "Catálogo plan validado", "CARRERA_PLAN", "CODCARR+ASIGNATURA", "CODIGO_UNICO+PLAN_ESTUDIOS+unidad", "N:1 si catálogo existe", "UNIDADES_CURSADAS/APROBADAS", "ALTO"),
        ("CRUCE-005", "Separar convalidaciones", "PROMEDIOS Hoja1", "PERSONA_ASIGNATURA", "Regla instructivo TXT", "REGLA", "CONVALIDADO+ESTADO", "NO_APLICA", "N:1", "marcas anual/acumulado", "ALTO"),
    ]
    out = [
        {
            "ID_CRUCE": rid,
            "OBJETIVO": objetivo,
            "ARCHIVO_IZQUIERDO": ai,
            "GRANULARIDAD_IZQUIERDA": gi,
            "ARCHIVO_DERECHO": ad,
            "GRANULARIDAD_DERECHA": gd,
            "LLAVE_IZQUIERDA": li,
            "LLAVE_DERECHA": ld,
            "CARDINALIDAD_ESPERADA": card,
            "COLUMNAS_INCORPORADAS": cols,
            "RIESGO_MANY_TO_MANY": riesgo,
            "VALIDACION_PREVIA": "Validar unicidad, cobertura, período y ausencia de datos personales en salidas versionables.",
            "VALIDACION_POSTERIOR": "Conteos, duplicados, no modificables, trazabilidad y auditoría de rechazos.",
            "ESTADO": "PROPUESTO_NO_EJECUTADO",
            "OBSERVACIONES": "No ejecutar hasta cerrar B03.",
        }
        for rid, objetivo, ai, gi, ad, gd, li, ld, card, cols, riesgo in rows
    ]
    write_tsv(ROOT / "07_control" / "MATRIZ_CRUCES_PROPUESTA.tsv", out, [
        "ID_CRUCE",
        "OBJETIVO",
        "ARCHIVO_IZQUIERDO",
        "GRANULARIDAD_IZQUIERDA",
        "ARCHIVO_DERECHO",
        "GRANULARIDAD_DERECHA",
        "LLAVE_IZQUIERDA",
        "LLAVE_DERECHA",
        "CARDINALIDAD_ESPERADA",
        "COLUMNAS_INCORPORADAS",
        "RIESGO_MANY_TO_MANY",
        "VALIDACION_PREVIA",
        "VALIDACION_POSTERIOR",
        "ESTADO",
        "OBSERVACIONES",
    ])


def write_pipeline_docs() -> None:
    carreras = """# Diseño Pipeline Carreras Avance Curricular 2026

No implementado productivamente.

1. Cargar precarga congelada.
2. Validar hash contra manifiesto.
3. Validar estructura y primera columna institucional.
4. Preservar campos no modificables.
5. Cargar fuente institucional de planes/mallas.
6. Resolver planes con evidencia.
7. Resolver unidad de medida.
8. Incorporar total de unidades.
9. Incorporar distribución anual.
10. Validar suma anual contra total.
11. Validar duración y años informados.
12. Resolver vigencia con evidencia.
13. Generar auditoría.
14. Generar archivo de control con encabezado.
15. Gate de aprobación.
16. Congelar catálogo de planes validado.

Gates: fuente institucional identificada, unidad no pendiente, total no pendiente, distribución conciliada, no modificables preservados, sin referencias activas al PDF retirado.
"""
    matricula = """# Diseño Pipeline Matrícula Avance Curricular 2026

No implementado productivamente.

1. Cargar precarga congelada.
2. Validar hash contra manifiesto.
3. Validar estructura y primera columna institucional.
4. Preservar campos no modificables.
5. Cargar catálogo validado de Carreras.
6. Asignar PLAN_ESTUDIOS solo con evidencia.
7. Cargar actividad 2025-1.
8. Cargar actividad 2025-2.
9. Calcular anual cursado.
10. Calcular anual aprobado.
11. Calcular acumulado cursado.
12. Calcular acumulado aprobado.
13. Tratar convalidaciones, validación de estudios y RAP explícitamente.
14. Validar unidad de medida.
15. Validar contra total del plan.
16. Aplicar control de tolerancia 25%.
17. Resolver vigencia con evidencia.
18. Clasificar duplicados.
19. Generar auditoría.
20. Generar archivo de control con encabezado.
21. Gate de aprobación.

Gates: Carreras validada, plan existente, unidad consistente, anual/acumulado separado, no se reemplazan desconocidos por cero.
"""
    gates = """# Orquestación y Gates

Orden obligatorio:

1. Validar carga congelada rectificada.
2. Procesar Carreras.
3. Validar Carreras.
4. Congelar catálogo de planes.
5. Procesar Matrícula.
6. Validar Matrícula.
7. Validar cruces.
8. Generar controles.
9. Generar PES solo con cero bloqueos.

Bloqueos productivos: B01 para commit, B03 para fuentes institucionales, B04 para privacidad, B05 para PES. B02 queda descartado por decisión interna DI-AC-2026-001.
"""
    plan = """# Plan de Implementación

| FASE | OBJETIVO | ENTRADAS | SALIDAS | DEPENDENCIAS | SCRIPT_PROPUESTO | PRUEBAS | GATES | BLOQUEOS | REANUDACION | ESTADO |
| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |
| 00A-R | Validar carga rectificada | Manifiestos | Estado rectificado | Ninguna | validar_carga_congelada.py | hashes | 4 fuentes activas | B04 | revalidar hashes | COMPLETADA |
| 04 | Gobernanza por columna | Matriz normativa, contratos | TSV por columna | 00A-R | no aplica | conteo 42 campos | sin PDF activo | B03 | regenerar gobernanza | COMPLETADA |
| 04B | Catálogo de validaciones | Instructivo TXT | catálogos TSV | 04 | no aplica | conteo validaciones | reglas separadas | B03 | regenerar catálogos | COMPLETADA |
| 04D | Perfil institucional | PROMEDIOS congelado | perfiles TSV | 00A-R | perfilar_fuentes_institucionales.py | sin muestras reales | privacidad | B03/B04 | re perfilar | COMPLETADA |
| 05-CAR | Implementar Carreras | Precarga + fuentes planes | control Carreras | B03 cerrado | procesar_carreras.py; validar_carreras.py; congelar_catalogo_planes.py | sintéticas | unidad/total/distribución | B03/B05 | reanudar por manifest | PENDIENTE |
| 05-MAT | Implementar Matrícula | Precarga + catálogo Carreras + historial | control Matrícula | Carreras validada | procesar_matricula.py; validar_matricula.py; validar_cruces.py | sintéticas | cruces y acumulados | B03/B05 | reanudar por manifest | PENDIENTE |
| 11-12 | Control y PES | controles aprobados | control/PES | gates cero | generar_control.py; generar_pes.py; auditoria_final.py | end-to-end | cero bloqueos | B05 | manifest ejecución | PENDIENTE |

No se crean scripts productivos en esta fase.
"""
    (ROOT / "01_documentacion" / "06_DISENO_PIPELINE_CARRERAS.md").write_text(carreras, encoding="utf-8")
    (ROOT / "01_documentacion" / "07_DISENO_PIPELINE_MATRICULA.md").write_text(matricula, encoding="utf-8")
    (ROOT / "01_documentacion" / "08_ORQUESTACION_Y_GATES.md").write_text(gates, encoding="utf-8")
    (ROOT / "01_documentacion" / "09_PLAN_IMPLEMENTACION.md").write_text(plan, encoding="utf-8")


def write_bloqueos(retiro_time: str) -> None:
    rows = [
        {
            "ID_BLOQUEO": "B01",
            "DESCRIPCION": "Working tree con entradas modificadas o no rastreadas de múltiples procesos o de esta ejecución.",
            "ORIGEN": "Git status",
            "FECHA_APERTURA": "2026-06-26",
            "ESTADO": "ABIERTO",
            "FECHA_CIERRE": "NO_APLICA",
            "DECISION_INTERNA": "NO_APLICA",
            "IMPACTO": "Impide commit indiscriminado.",
            "ACCION_REQUERIDA": "Usar CONCILIACION_WORKING_TREE.tsv antes de cualquier commit selectivo.",
            "EVIDENCIA": "avance_curricular_2026/07_control/CONCILIACION_WORKING_TREE.tsv",
            "OBSERVACIONES": "No se hizo commit.",
        },
        {
            "ID_BLOQUEO": "B02",
            "DESCRIPCION": "PDF no extraído textualmente.",
            "ORIGEN": "Ejecución anterior",
            "FECHA_APERTURA": "2026-06-26",
            "ESTADO": "DESCARTADO_POR_DECISION_INTERNA",
            "FECHA_CIERRE": retiro_time,
            "DECISION_INTERNA": DECISION_ID,
            "IMPACTO": "No bloquea ninguna fase.",
            "ACCION_REQUERIDA": "Ninguna; no se requiere extracción ni comparación.",
            "EVIDENCIA": "avance_curricular_2026/01_documentacion/DECISION_INTERNA_DI-AC-2026-001.md",
            "OBSERVACIONES": "TXT es fuente normativa única.",
        },
        {
            "ID_BLOQUEO": "B03",
            "DESCRIPCION": "Faltan fuentes institucionales definitivas para planes, unidades, distribución y avance.",
            "ORIGEN": "Matriz de fuentes",
            "FECHA_APERTURA": "2026-06-26",
            "ESTADO": "ABIERTO",
            "FECHA_CIERRE": "NO_APLICA",
            "DECISION_INTERNA": "NO_APLICA",
            "IMPACTO": "Bloquea implementación productiva y PES.",
            "ACCION_REQUERIDA": "Aportar/validar fuentes institucionales y llaves.",
            "EVIDENCIA": "avance_curricular_2026/09_pendientes/FUENTES_INSTITUCIONALES_FALTANTES.tsv",
            "OBSERVACIONES": "PROMEDIOSDEALUMNOS es candidato parcial, no fuente suficiente automática.",
        },
        {
            "ID_BLOQUEO": "B04",
            "DESCRIPCION": "Repositorio público con datos personales locales.",
            "ORIGEN": "Visibilidad GitHub",
            "FECHA_APERTURA": "2026-06-26",
            "ESTADO": "CONTROLADO_ABIERTO",
            "FECHA_CIERRE": "NO_APLICA",
            "DECISION_INTERNA": "NO_APLICA",
            "IMPACTO": "Originales y derivados personales fuera de Git.",
            "ACCION_REQUERIDA": "Mantener .gitignore y no añadir datos personales.",
            "EVIDENCIA": "avance_curricular_2026/01_documentacion/POLITICA_DATOS_Y_GIT.md",
            "OBSERVACIONES": "No se incluyeron muestras reales en artefactos versionables.",
        },
        {
            "ID_BLOQUEO": "B05",
            "DESCRIPCION": "No corresponde generar PES mientras falten gobernanza, linaje, fuentes y validaciones.",
            "ORIGEN": "Gate de proceso",
            "FECHA_APERTURA": "2026-06-26",
            "ESTADO": "ABIERTO",
            "FECHA_CIERRE": "NO_APLICA",
            "DECISION_INTERNA": "NO_APLICA",
            "IMPACTO": "Bloquea salidas PES.",
            "ACCION_REQUERIDA": "Cerrar B03 y aprobar gates antes de PES.",
            "EVIDENCIA": "avance_curricular_2026/01_documentacion/08_ORQUESTACION_Y_GATES.md",
            "OBSERVACIONES": "No se generó PES.",
        },
    ]
    write_tsv(ROOT / "09_pendientes" / "REGISTRO_BLOQUEOS.tsv", rows, [
        "ID_BLOQUEO",
        "DESCRIPCION",
        "ORIGEN",
        "FECHA_APERTURA",
        "ESTADO",
        "FECHA_CIERRE",
        "DECISION_INTERNA",
        "IMPACTO",
        "ACCION_REQUERIDA",
        "EVIDENCIA",
        "OBSERVACIONES",
    ])


def write_conciliation() -> dict[str, Counter]:
    status_lines = run(["git", "status", "--porcelain=v1", "-uall"]).splitlines()
    rows = []
    for line in status_lines:
        if not line:
            continue
        status = line[:2].strip() or line[:2]
        path = line[3:]
        if " -> " in path:
            path = path.split(" -> ", 1)[1]
        clean_path = mask_inventory(path.strip('"'))
        category, personal, action = classify_path(path.strip('"'))
        generated_00 = "SI" if category == "GENERADO_FASE_00A_03" else "NO"
        generated_ret = "SI" if category == "GENERADO_RECTIFICACION" or (path.endswith(PDF_NAME) or "COMPARACION_INSTRUCTIVO" in path) else "NO"
        generated_04 = "SI" if category == "GENERADO_FASE_04" else "NO"
        rows.append(
            {
                "RUTA": clean_path,
                "ESTADO_GIT": status,
                "EXISTIA_ANTES_FASE_00A": "NO_DETERMINADO",
                "GENERADO_EN_FASE_00A_03": generated_00,
                "GENERADO_EN_RECTIFICACION": generated_ret,
                "GENERADO_EN_FASE_04": generated_04,
                "PROCESO_ASOCIADO": "Avance Curricular SIES 2026" if path.startswith("avance_curricular_2026/") else "OTRO_PROCESO_O_PREEXISTENTE",
                "CONTIENE_DATOS_PERSONALES": personal,
                "IGNORADO_POR_GIT": git_check_ignore(path.strip('"')),
                "APTO_PARA_GIT": "NO" if personal == "SI" or path.startswith("avance_curricular_2026/00_fuentes_congeladas") else ("SI_CON_REVISION" if action == "CANDIDATO_COMMIT_POSTERIOR" else "NO"),
                "ACCION_RECOMENDADA": action,
                "JUSTIFICACION": "Conciliación documental; no hacer commit automático.",
                "PENDIENTE": "REVISAR_ANTES_DE_COMMIT" if action != "IGNORAR_GIT" else "MANTENER_LOCAL",
            }
        )
    fields = [
        "RUTA",
        "ESTADO_GIT",
        "EXISTIA_ANTES_FASE_00A",
        "GENERADO_EN_FASE_00A_03",
        "GENERADO_EN_RECTIFICACION",
        "GENERADO_EN_FASE_04",
        "PROCESO_ASOCIADO",
        "CONTIENE_DATOS_PERSONALES",
        "IGNORADO_POR_GIT",
        "APTO_PARA_GIT",
        "ACCION_RECOMENDADA",
        "JUSTIFICACION",
        "PENDIENTE",
    ]
    write_tsv(ROOT / "07_control" / "CONCILIACION_WORKING_TREE.tsv", rows, fields)
    counts = Counter(row["ACCION_RECOMENDADA"] for row in rows)
    personal_count = sum(1 for row in rows if row["CONTIENE_DATOS_PERSONALES"] == "SI")
    md = f"""# Conciliación Git

Fecha: {NOW.isoformat()} America/Santiago.

## Resumen

- Total de entradas: {len(rows)}
- Datos personales locales: {personal_count}
- Candidatas a commit posterior: {counts.get('CANDIDATO_COMMIT_POSTERIOR', 0)}
- Conservar sin commit/ignorar: {counts.get('IGNORAR_GIT', 0)}
- No tocar/revisar manualmente: {counts.get('NO_TOCAR', 0) + counts.get('REVISAR_MANUALMENTE', 0)}

## Lectura

Los cambios de Avance Curricular pertenecen a la rectificación, gobernanza, catálogos, perfiles, matrices y diseño documental. Las fuentes congeladas con datos personales no son aptas para Git en un repositorio público.

El retiro controlado del PDF aparece como eliminación dentro de `00_fuentes_congeladas`; la trazabilidad queda en la auditoría y el manifiesto.

Es posible construir un commit selectivo futuro solo con documentación saneada, contratos, gobernanza, catálogos, scripts documentales y controles sin datos personales. No debe hacerse commit indiscriminado.

B01 permanece ABIERTO hasta ejecutar revisión selectiva.
"""
    (ROOT / "01_documentacion" / "04_CONCILIACION_GIT.md").write_text(md, encoding="utf-8")
    return {"counts": counts, "personal": Counter({"SI": personal_count, "NO": len(rows) - personal_count}), "total": Counter({"entries": len(rows)})}


def final_control(rect: dict[str, Any]) -> dict[str, Any]:
    checks = {
        "branch": run(["git", "branch", "--show-current"]),
        "commit": run(["git", "rev-parse", "--short", "HEAD"]),
        "pdf_in_freeze_exists": any(FREEZE.rglob(PDF_NAME)),
        "downloads_pdf_exists": DOWNLOADS_PDF.exists(),
        "txt_hash": sha256(ORIG / TXT_NAME),
        "carreras_hash": sha256(ORIG / "5810_Precarga Carreras Avance Curricular 20268.csv"),
        "matricula_hash": sha256(next(ORIG.glob("5809_Precarga Matri*cula Avance Curricular 2026.csv"))),
        "promedios_hash": sha256(ORIG / PROMEDIOS_NAME),
        "pes_files": list((ROOT / "11_archivos_subida").glob("*PES*")) if (ROOT / "11_archivos_subida").exists() else [],
    }
    if checks["branch"] != "backup/pre-sync-fix-20260410-avance":
        raise RuntimeError("Rama incorrecta")
    if checks["pdf_in_freeze_exists"]:
        raise RuntimeError("El PDF aún existe en la carga congelada")
    if not checks["downloads_pdf_exists"]:
        raise RuntimeError("El PDF original de Descargas no está conservado")
    if checks["txt_hash"] != rect["txt_hash_before"] or checks["txt_hash"] != rect["txt_hash_after"]:
        raise RuntimeError("Hash TXT no coincide")
    if checks["pes_files"]:
        raise RuntimeError("Se detectaron archivos PES generados")
    write_json(ROOT / "07_control" / "CONTROL_FINAL_FASE_00AR_05P.json", checks)
    return checks


def main() -> int:
    ensure_dirs()
    branch = run(["git", "branch", "--show-current"])
    if branch != "backup/pre-sync-fix-20260410-avance":
        raise RuntimeError(f"Rama activa incorrecta: {branch}")
    rect = phase_00ar()
    gov = phase_04_governance()
    vals = phase_04b_validations()
    prom = phase_04d_promedios()
    sources = phase_04e_sources()
    write_pipeline_docs()
    write_bloqueos(rect["retiro_time"])
    conciliation = write_conciliation()
    checks = final_control(rect)
    summary = {
        "rectificacion": rect,
        "gobernanza": {k: dict(v) for k, v in gov.items()},
        "validaciones": {k: dict(v) for k, v in vals.items()},
        "promedios": prom,
        "fuentes": sources,
        "conciliacion": {
            "counts": dict(conciliation["counts"]),
            "personal": dict(conciliation["personal"]),
            "total": dict(conciliation["total"]),
        },
        "control_final": checks,
    }
    write_json(ROOT / "07_control" / "RESUMEN_FASE_00AR_05P.json", summary)
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
