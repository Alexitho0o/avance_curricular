from pathlib import Path
from datetime import datetime
import pandas as pd
import numpy as np
import json
import hashlib
import shutil
import zipfile
import re

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# =============================================================================
# CONTEXTO
# =============================================================================

PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Mapeo resolución bloqueos 5809 lógica 2025"
ANIO_REFERENCIA = 2025

RAIZ = Path("/Users/alexi/Documents/GitHub/avance_curricular").resolve()
DESKTOP = Path.home() / "Desktop"

AUDITORIA_CORREGIDA = Path(
    "/Users/alexi/Documents/GitHub/avance_curricular/avance_curricular_2026/"
    "29_auditoria_corregida_logica_2025_5809/"
    "AUDITORIA_CORREGIDA_LOGICA_2025_5809.xlsx"
)

PROMEDIOS = Path(
    "/Users/alexi/Documents/GitHub/avance_curricular/avance_curricular_2026/"
    "25_actualizacion_fuente_promedios/"
    "CAMBIO_PROMEDIOSDEALUMNOS_20260704_221059/"
    "00_FUENTE_CONGELADA/"
    "PROMEDIOSDEALUMNOS_7804_ACTUALIZADO_20260704_221059.xlsx"
)

CSV_5809 = Path(
    "/Users/alexi/Documents/GitHub/avance_curricular/avance_curricular_2026/"
    "14_sies_ready_5809/"
    "SIES_READY_TECNICO_5809_22_COLUMNAS_20260704_151418/"
    "01_ARCHIVO_SIES_READY_TECNICO/"
    "5809_MATRICULA_AVANCE_CURRICULAR_2026_SIES_READY_TECNICO.csv"
)

timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

SALIDA = (
    RAIZ
    / "avance_curricular_2026"
    / "30_mapeo_resolucion_bloqueos_5809_2025"
    / f"MAPEO_RESOLUCION_BLOQUEOS_5809_2025_{timestamp}"
)

RESULTADOS = SALIDA / "02_RESULTADOS"
AUDITORIA = SALIDA / "03_AUDITORIA"
REPORTES = SALIDA / "04_REPORTES"
MANIFEST_DIR = SALIDA / "05_MANIFEST"

for carpeta in [RESULTADOS, AUDITORIA, REPORTES, MANIFEST_DIR]:
    carpeta.mkdir(parents=True, exist_ok=True)

CARPETA_ESCRITORIO = (
    DESKTOP
    / f"AVANCE_CURRICULAR_2026_MAPEO_RESOLUCION_BLOQUEOS_5809_2025_{timestamp}"
)

CARPETA_ESCRITORIO.mkdir(parents=True, exist_ok=True)


# =============================================================================
# UTILIDADES
# =============================================================================

def sha256(path):
    if not path.exists():
        return "NO_EXISTE"

    h = hashlib.sha256()

    with path.open("rb") as f:
        for bloque in iter(lambda: f.read(1024 * 1024), b""):
            h.update(bloque)

    return h.hexdigest()


def json_safe(obj):
    if isinstance(obj, Path):
        return str(obj)

    if isinstance(obj, np.integer):
        return int(obj)

    if isinstance(obj, np.floating):
        return float(obj)

    if isinstance(obj, np.bool_):
        return bool(obj)

    return obj


def dumps_json_safe(obj):
    return json.dumps(
        obj,
        ensure_ascii=False,
        indent=2,
        default=json_safe,
    )


def norm_text(x):
    txt = str(x or "").strip().upper()

    reemplazos = {
        "Á": "A",
        "É": "E",
        "Í": "I",
        "Ó": "O",
        "Ú": "U",
        "Ñ": "N",
    }

    for a, b in reemplazos.items():
        txt = txt.replace(a, b)

    txt = re.sub(r"\s+", " ", txt)
    return txt.strip()


def norm_col(x):
    txt = norm_text(x)
    txt = re.sub(r"[^A-Z0-9]+", "_", txt)
    txt = re.sub(r"_+", "_", txt)
    return txt.strip("_")


def to_bool(x):
    if isinstance(x, bool):
        return x

    txt = norm_text(x)

    if txt in ["TRUE", "VERDADERO", "SI", "SÍ", "1", "OK"]:
        return True

    return False


def to_number(x):
    if x is None:
        return 0

    try:
        return float(str(x).replace(",", "."))
    except Exception:
        return 0


def normalizar_rut_libre(x):
    txt = str(x or "").strip().upper()
    txt = txt.replace(".", "").replace(" ", "")

    if "-" in txt:
        cuerpo, dv = txt.split("-", 1)
        cuerpo = re.sub(r"\D", "", cuerpo)
        dv = re.sub(r"[^0-9K]", "", dv)

        if cuerpo and dv:
            return f"{cuerpo}-{dv[-1]}"

        return ""

    limpio = re.sub(r"[^0-9K]", "", txt)

    if len(limpio) >= 2:
        return f"{limpio[:-1]}-{limpio[-1]}"

    return ""


def rut_sin_dv(x):
    rut = normalizar_rut_libre(x)

    if "-" in rut:
        return rut.split("-", 1)[0]

    return re.sub(r"\D", "", str(x or ""))


def read_sheet_or_empty(path, sheet):
    try:
        return pd.read_excel(
            path,
            sheet_name=sheet,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl",
        )
    except Exception:
        return pd.DataFrame()


def detectar_columna(df, candidatos):
    if df.empty:
        return None

    mapa = {norm_col(c): c for c in df.columns}

    for cand in candidatos:
        key = norm_col(cand)

        if key in mapa:
            return mapa[key]

    for c in df.columns:
        nc = norm_col(c)

        for cand in candidatos:
            if norm_col(cand) in nc:
                return c

    return None


def agregar_columna_si_no_existe(df, columna, valor=""):
    if columna not in df.columns:
        df[columna] = valor

    return df


def validar_xlsx(path):
    with zipfile.ZipFile(path, "r") as zf:
        corrupto = zf.testzip()

        if corrupto is not None:
            raise RuntimeError(f"XLSX corrupto en {corrupto}")

    wb = load_workbook(path, read_only=True)
    hojas = wb.sheetnames
    wb.close()

    return hojas


def aplicar_formato_excel(path):
    wb = load_workbook(path)

    fill_header = PatternFill("solid", fgColor="1F4E78")
    fill_ok = PatternFill("solid", fgColor="C6EFCE")
    fill_warn = PatternFill("solid", fgColor="FFEB9C")
    fill_bad = PatternFill("solid", fgColor="FFC7CE")
    fill_info = PatternFill("solid", fgColor="D9EAF7")

    font_header = Font(color="FFFFFF", bold=True)
    font_ok = Font(color="006100", bold=True)
    font_warn = Font(color="9C6500", bold=True)
    font_bad = Font(color="9C0006", bold=True)

    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    align_center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    for ws in wb.worksheets:
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = fill_header
                cell.font = font_header
                cell.alignment = align_center
                cell.border = border

        for row in ws.iter_rows(
            min_row=1,
            max_row=min(ws.max_row, 3000),
            min_col=1,
            max_col=ws.max_column,
        ):
            for cell in row:
                cell.alignment = align_center
                cell.border = border

        ws.freeze_panes = "A2"

        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)

            max_len = 0

            for row_idx in range(1, min(ws.max_row, 500) + 1):
                value = ws.cell(row_idx, col_idx).value

                if value is not None:
                    max_len = max(max_len, len(str(value)))

            header = str(ws.cell(1, col_idx).value or "").upper()

            width = min(max(max_len + 2, 12), 70)

            if any(
                token in header
                for token in [
                    "RUTA",
                    "ACCION",
                    "MOTIVO",
                    "OBSERVACION",
                    "CRITERIO",
                    "DETALLE",
                    "COMO",
                    "JUSTIFICACION",
                ]
            ):
                width = 70

            ws.column_dimensions[col_letter].width = width

        for row_idx in range(2, min(ws.max_row, 3000) + 1):
            row_text = " ".join(
                str(ws.cell(row_idx, c).value or "").upper()
                for c in range(1, ws.max_column + 1)
            )

            if "BLOQUEO" in row_text or "NO_LISTO" in row_text:
                fill = fill_bad
                font = font_bad
            elif "REVISAR" in row_text or "DIFERENCIA" in row_text or "PENDIENTE" in row_text:
                fill = fill_warn
                font = font_warn
            elif "OK" in row_text:
                fill = fill_ok
                font = font_ok
            else:
                fill = None
                font = None

            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row_idx, c).fill = fill
                    ws.cell(row_idx, c).font = font

        if ws.title.startswith("P1") or ws.title.startswith("P2") or ws.title.startswith("P3") or ws.title.startswith("P4"):
            ws.sheet_properties.tabColor = "FFC000"
        elif "DICTAMEN" in ws.title or "RESUMEN" in ws.title:
            ws.sheet_properties.tabColor = "5B9BD5"
        elif "FUENTES" in ws.title or "VALIDACIONES" in ws.title:
            ws.sheet_properties.tabColor = "7030A0"
        else:
            ws.sheet_properties.tabColor = "70AD47"

    wb.save(path)


# =============================================================================
# VALIDACIÓN DE ENTRADAS
# =============================================================================

bloqueos = []

if not AUDITORIA_CORREGIDA.exists():
    bloqueos.append(f"No existe auditoría corregida: {AUDITORIA_CORREGIDA}")

if not PROMEDIOS.exists():
    bloqueos.append(f"No existe PROMEDIOS: {PROMEDIOS}")

if not CSV_5809.exists():
    bloqueos.append(f"No existe CSV 5809: {CSV_5809}")

if bloqueos:
    reporte_bloqueo = REPORTES / "BLOQUEO_ENTRADAS_MAPEO_RESOLUCION.md"

    reporte_bloqueo.write_text(
        "# Bloqueo de entradas\n\n"
        + "\n".join(f"- {b}" for b in bloqueos)
        + "\n",
        encoding="utf-8",
    )

    print()
    print("=" * 120)
    print("BLOQUEO DE ENTRADAS")
    print("=" * 120)

    for b in bloqueos:
        print(f"- {b}")

    print(f"Reporte: {reporte_bloqueo}")
    print("=" * 120)
    raise SystemExit(1)


# =============================================================================
# CARGA AUDITORÍA PREVIA
# =============================================================================

dictamen_prev = read_sheet_or_empty(AUDITORIA_CORREGIDA, "DICTAMEN_GLOBAL")
validaciones_prev = read_sheet_or_empty(AUDITORIA_CORREGIDA, "VALIDACIONES")
columnas_prev = read_sheet_or_empty(AUDITORIA_CORREGIDA, "COLUMNAS_DETECTADAS")
catalogo_estado = read_sheet_or_empty(AUDITORIA_CORREGIDA, "CATALOGO_ESTADO")
mapeo_codcli_rut = read_sheet_or_empty(AUDITORIA_CORREGIDA, "MAPEO_CODCLI_RUT")
aud_anual = read_sheet_or_empty(AUDITORIA_CORREGIDA, "AUDITORIA_ANUAL_2025_FILA_A_FILA")
filas_revisar_anual = read_sheet_or_empty(AUDITORIA_CORREGIDA, "FILAS_REVISAR_ANUAL_2025")
resumen_anual = read_sheet_or_empty(AUDITORIA_CORREGIDA, "RESUMEN_ANUAL_2025")
aud_acum = read_sheet_or_empty(AUDITORIA_CORREGIDA, "AUDITORIA_ACUMULADO_2025_FILA_A_FILA")
filas_revisar_acum = read_sheet_or_empty(AUDITORIA_CORREGIDA, "FILAS_REVISAR_ACUMULADO_2025")
resumen_acum = read_sheet_or_empty(AUDITORIA_CORREGIDA, "RESUMEN_ACUMULADO_2025")
logica = read_sheet_or_empty(AUDITORIA_CORREGIDA, "LOGICA_AUDITADA")
fuentes_prev = read_sheet_or_empty(AUDITORIA_CORREGIDA, "FUENTES")

hojas_requeridas = [
    ("DICTAMEN_GLOBAL", dictamen_prev),
    ("VALIDACIONES", validaciones_prev),
    ("COLUMNAS_DETECTADAS", columnas_prev),
    ("CATALOGO_ESTADO", catalogo_estado),
    ("MAPEO_CODCLI_RUT", mapeo_codcli_rut),
    ("AUDITORIA_ANUAL_2025_FILA_A_FILA", aud_anual),
    ("FILAS_REVISAR_ANUAL_2025", filas_revisar_anual),
    ("RESUMEN_ANUAL_2025", resumen_anual),
    ("AUDITORIA_ACUMULADO_2025_FILA_A_FILA", aud_acum),
    ("FILAS_REVISAR_ACUMULADO_2025", filas_revisar_acum),
    ("RESUMEN_ACUMULADO_2025", resumen_acum),
    ("LOGICA_AUDITADA", logica),
    ("FUENTES", fuentes_prev),
]

validacion_hojas = []

for nombre, df in hojas_requeridas:
    validacion_hojas.append({
        "HOJA": nombre,
        "EXISTE_Y_TIENE_FILAS": "SI" if not df.empty else "NO",
        "FILAS": len(df),
        "COLUMNAS": len(df.columns),
    })

validacion_hojas_df = pd.DataFrame(validacion_hojas)

if aud_anual.empty:
    raise SystemExit("BLOQUEO: no se pudo leer AUDITORIA_ANUAL_2025_FILA_A_FILA.")

if aud_acum.empty:
    raise SystemExit("BLOQUEO: no se pudo leer AUDITORIA_ACUMULADO_2025_FILA_A_FILA.")


# =============================================================================
# CARGA PROMEDIOS PARA ENRIQUECER CANDIDATOS
# =============================================================================

promedios_hojas = {}

try:
    xls = pd.ExcelFile(PROMEDIOS, engine="openpyxl")

    for hoja in xls.sheet_names:
        promedios_hojas[hoja] = pd.read_excel(
            PROMEDIOS,
            sheet_name=hoja,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl",
        )
except Exception:
    promedios_hojas = {}


def construir_mapeo_desde_promedios():
    filas = []

    for hoja, df in promedios_hojas.items():
        col_codcli = detectar_columna(df, ["CODCLI", "COD_CLI", "CODIGO_CLIENTE", "CODIGO_ALUMNO"])
        col_rut = detectar_columna(df, ["RUT", "RUN", "NUM_DOCUMENTO", "NUMDOCUMENTO", "DOCUMENTO"])
        col_dv = detectar_columna(df, ["DV", "DIGITO_VERIFICADOR", "DVRUT"])
        col_codigo_unico = detectar_columna(df, ["CODIGO_UNICO", "COD_UNICO"])

        cols_nombre = [
            c for c in df.columns
            if any(token in norm_col(c) for token in ["NOMBRE", "APELLIDO"])
        ]

        if not col_codcli:
            continue

        for _, r in df.iterrows():
            codcli = str(r.get(col_codcli, "")).strip()

            if not codcli:
                continue

            if col_rut and col_dv and col_rut != col_dv:
                rut = normalizar_rut_libre(f"{r.get(col_rut, '')}-{r.get(col_dv, '')}")
            elif col_rut:
                rut = normalizar_rut_libre(r.get(col_rut, ""))
            else:
                rut = ""

            codigo_unico = str(r.get(col_codigo_unico, "")).strip() if col_codigo_unico else ""

            if cols_nombre:
                nombre = norm_text(" ".join(str(r.get(c, "")) for c in cols_nombre))
            else:
                nombre = ""

            if not rut and not codigo_unico and not nombre:
                continue

            filas.append({
                "HOJA_ORIGEN": hoja,
                "CODCLI": codcli,
                "RUT_NORMALIZADO": rut,
                "RUT_SIN_DV": rut_sin_dv(rut),
                "CODIGO_UNICO": codigo_unico,
                "NOMBRE_COMPUESTO": nombre,
            })

    if not filas:
        return pd.DataFrame(
            columns=[
                "HOJA_ORIGEN",
                "CODCLI",
                "RUT_NORMALIZADO",
                "RUT_SIN_DV",
                "CODIGO_UNICO",
                "NOMBRE_COMPUESTO",
            ]
        )

    return pd.DataFrame(filas).drop_duplicates()


mapeo_promedios = construir_mapeo_desde_promedios()


# =============================================================================
# NORMALIZACIÓN AUDITORÍA ANUAL
# =============================================================================

for col in [
    "RUT_NORMALIZADO",
    "CODCLI_LISTA",
    "CODCLI_CANTIDAD",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
    "DICTAMEN_ANUAL_2025",
]:
    aud_anual = agregar_columna_si_no_existe(aud_anual, col, "")

for col in [
    "OK_CURSO_1ER_SEM",
    "OK_CURSO_2DO_SEM",
    "OK_UNIDADES_CURSADAS_2025",
    "OK_UNIDADES_APROBADAS_2025",
    "TIENE_EVIDENCIA_2025",
    "TIENE_EVIDENCIA_2026_MISMO_CODCLI",
]:
    aud_anual = agregar_columna_si_no_existe(aud_anual, col, "False")

for col in [
    "DIF_UNIDADES_CURSADAS_2025",
    "DIF_UNIDADES_APROBADAS_2025",
    "REGISTROS_2025",
    "REGISTROS_2026_MISMO_CODCLI",
]:
    aud_anual = agregar_columna_si_no_existe(aud_anual, col, 0)

aud_anual["CODCLI_CANTIDAD_NUM"] = aud_anual["CODCLI_CANTIDAD"].apply(to_number)
aud_anual["RUT_SIN_DV"] = aud_anual["RUT_NORMALIZADO"].apply(rut_sin_dv)

if "NOMBRES" in aud_anual.columns:
    aud_anual["NOMBRE_5809_BASE"] = aud_anual["NOMBRES"].apply(norm_text)
else:
    aud_anual["NOMBRE_5809_BASE"] = ""


# =============================================================================
# P1 — SIN CODCLI
# =============================================================================

p1_sin_codcli = aud_anual[
    aud_anual["DICTAMEN_ANUAL_2025"].astype(str).str.contains("SIN_CODCLI", case=False, na=False)
    | aud_anual["CODCLI_CANTIDAD_NUM"].eq(0)
].copy()

candidatos = []

for _, row in p1_sin_codcli.iterrows():
    rut = str(row.get("RUT_NORMALIZADO", "")).strip()
    rut_num = rut_sin_dv(rut)
    codigo_unico = str(row.get("CODIGO_UNICO", "")).strip()
    nombre = norm_text(row.get("NOMBRE_5809_BASE", ""))

    if mapeo_promedios.empty:
        candidatos.append({
            "RUT_5809": rut,
            "CODIGO_UNICO_5809": codigo_unico,
            "CANDIDATO_CODCLI": "",
            "SCORE": 0,
            "ACCION_PROPUESTA": "SOLICITAR_CODCLI_A_DOCENCIA_O_BASE_MAESTRA",
            "MOTIVO": "No existe mapeo disponible desde PROMEDIOS.",
        })
        continue

    cand = mapeo_promedios.copy()

    cand["MATCH_RUT_COMPLETO"] = cand["RUT_NORMALIZADO"].eq(rut) & cand["RUT_NORMALIZADO"].ne("")
    cand["MATCH_RUT_SIN_DV"] = cand["RUT_SIN_DV"].eq(rut_num) & cand["RUT_SIN_DV"].ne("")
    cand["MATCH_CODIGO_UNICO"] = cand["CODIGO_UNICO"].eq(codigo_unico) & cand["CODIGO_UNICO"].ne("")
    cand["MATCH_NOMBRE"] = cand["NOMBRE_COMPUESTO"].apply(
        lambda x: bool(x and nombre and (x in nombre or nombre in x))
    )

    cand["SCORE"] = (
        cand["MATCH_RUT_COMPLETO"].astype(int) * 100
        + cand["MATCH_RUT_SIN_DV"].astype(int) * 60
        + cand["MATCH_CODIGO_UNICO"].astype(int) * 40
        + cand["MATCH_NOMBRE"].astype(int) * 20
    )

    cand = cand[cand["SCORE"].gt(0)].copy()

    if cand.empty:
        candidatos.append({
            "RUT_5809": rut,
            "CODIGO_UNICO_5809": codigo_unico,
            "CANDIDATO_CODCLI": "",
            "SCORE": 0,
            "ACCION_PROPUESTA": "SOLICITAR_CODCLI_A_DOCENCIA_O_BASE_MAESTRA",
            "MOTIVO": "No se encontró candidato por RUT, RUT sin DV, código único ni nombre.",
        })
    else:
        cand = cand.sort_values(["SCORE", "HOJA_ORIGEN"], ascending=[False, True])

        for _, c in cand.head(5).iterrows():
            if c["SCORE"] >= 100:
                accion = "VALIDAR_MAPEO_POR_RUT_COMPLETO"
            elif c["SCORE"] >= 60:
                accion = "REVISAR_MAPEO_POR_RUT_SIN_DV"
            elif c["SCORE"] >= 40:
                accion = "REVISAR_MAPEO_POR_CODIGO_UNICO"
            else:
                accion = "REVISION_MANUAL_POR_NOMBRE"

            candidatos.append({
                "RUT_5809": rut,
                "CODIGO_UNICO_5809": codigo_unico,
                "PLAN_ESTUDIOS_5809": row.get("PLAN_ESTUDIOS", ""),
                "CANDIDATO_CODCLI": c["CODCLI"],
                "CANDIDATO_RUT": c["RUT_NORMALIZADO"],
                "CANDIDATO_CODIGO_UNICO": c["CODIGO_UNICO"],
                "CANDIDATO_NOMBRE": c["NOMBRE_COMPUESTO"],
                "HOJA_ORIGEN": c["HOJA_ORIGEN"],
                "SCORE": int(c["SCORE"]),
                "MATCH_RUT_COMPLETO": bool(c["MATCH_RUT_COMPLETO"]),
                "MATCH_RUT_SIN_DV": bool(c["MATCH_RUT_SIN_DV"]),
                "MATCH_CODIGO_UNICO": bool(c["MATCH_CODIGO_UNICO"]),
                "MATCH_NOMBRE": bool(c["MATCH_NOMBRE"]),
                "ACCION_PROPUESTA": accion,
                "MOTIVO": "Candidato de mapeo; validar antes de corregir lógica o archivo.",
            })

p1_candidatos = pd.DataFrame(candidatos)


# =============================================================================
# P2 — DIFERENCIAS 16-19
# =============================================================================

def row_bool(df, col):
    if col not in df.columns:
        return pd.Series([False] * len(df), index=df.index)

    return df[col].apply(to_bool)


mask_dif_16_19 = (
    (~row_bool(aud_anual, "OK_CURSO_1ER_SEM"))
    | (~row_bool(aud_anual, "OK_CURSO_2DO_SEM"))
    | (~row_bool(aud_anual, "OK_UNIDADES_CURSADAS_2025"))
    | (~row_bool(aud_anual, "OK_UNIDADES_APROBADAS_2025"))
    | aud_anual["DICTAMEN_ANUAL_2025"].astype(str).str.contains("DIFERENCIA", case=False, na=False)
)

p2_diferencias = aud_anual[mask_dif_16_19].copy()

def tipo_diferencia(row):
    tipos = []

    if not to_bool(row.get("OK_CURSO_1ER_SEM", False)):
        tipos.append("CURSO_1ER_SEM")

    if not to_bool(row.get("OK_CURSO_2DO_SEM", False)):
        tipos.append("CURSO_2DO_SEM")

    if not to_bool(row.get("OK_UNIDADES_CURSADAS_2025", False)):
        tipos.append("UNIDADES_CURSADAS")

    if not to_bool(row.get("OK_UNIDADES_APROBADAS_2025", False)):
        tipos.append("UNIDADES_APROBADAS")

    if not tipos:
        tipos.append("DIFERENCIA_NO_DETALLADA")

    return " | ".join(tipos)


def accion_diferencia(row):
    tipo = row.get("TIPO_DIFERENCIA_16_19", "")

    if "UNIDADES_APROBADAS" in tipo:
        return "REVISAR_ESTADO_APROBACION_Y_CATALOGO_ESTADO"

    if "UNIDADES_CURSADAS" in tipo:
        return "REVISAR_CONTEO_CODRAMO_DUPLICADOS_FILTRO_ANO_2025"

    if "CURSO" in tipo:
        return "REVISAR_PERIODO_2025_Y_EXISTENCIA_ASIGNATURAS"

    return "REVISION_MANUAL"

p2_diferencias["TIPO_DIFERENCIA_16_19"] = p2_diferencias.apply(tipo_diferencia, axis=1)
p2_diferencias["ACCION_PROPUESTA"] = p2_diferencias.apply(accion_diferencia, axis=1)


# =============================================================================
# P3 — ESTADOS NO CLASIFICADOS
# =============================================================================

if not catalogo_estado.empty:
    col_clasif = detectar_columna(catalogo_estado, ["CLASIFICACION_ESTADO", "CLASIFICACION"])
    col_estado = detectar_columna(catalogo_estado, ["ESTADO_AUDIT", "ESTADO"])
    col_freq = detectar_columna(catalogo_estado, ["FRECUENCIA_2025", "FRECUENCIA"])

    if col_clasif:
        p3_estados = catalogo_estado[
            catalogo_estado[col_clasif].astype(str).str.contains("NO_CLASIFICADO", case=False, na=False)
        ].copy()
    else:
        p3_estados = pd.DataFrame()
else:
    p3_estados = pd.DataFrame()

if p3_estados.empty:
    p3_estados = pd.DataFrame([
        {
            "ESTADO": "SIN_ESTADOS_NO_CLASIFICADOS_DETECTADOS_EN_CATALOGO",
            "ACCION_PROPUESTA": "SIN_ACCION",
        }
    ])
else:
    p3_estados["ACCION_PROPUESTA"] = (
        "SOLICITAR_CLASIFICACION_FUNCIONAL: APROBADO / NO_APROBADO. "
        "No clasificar automáticamente sin respaldo."
    )

p3_filas_impactadas = aud_anual[
    aud_anual["DICTAMEN_ANUAL_2025"].astype(str).str.contains("ESTADO", case=False, na=False)
].copy()


# =============================================================================
# P4 — ACUMULADO
# =============================================================================

for col in [
    "DICTAMEN_ACUMULADO_2025",
    "OK_ACUMULADO_B1",
    "RUT_NORMALIZADO",
    "CODCLI_LISTA",
    "CODIGO_UNICO",
    "PLAN_ESTUDIOS",
]:
    aud_acum = agregar_columna_si_no_existe(aud_acum, col, "")

mask_acum_revisar = (
    aud_acum["DICTAMEN_ACUMULADO_2025"].astype(str).str.contains("REVISAR|BLOQUEO|NO_EJECUTABLE", case=False, na=False)
    | (~aud_acum["OK_ACUMULADO_B1"].apply(to_bool))
)

p4_acum_revisar = aud_acum[mask_acum_revisar].copy()

p4_acum_revisar["ACCION_PROPUESTA"] = (
    "Revisar lógica acumulada separada de columnas 16-19. "
    "Validar si acumulado debe acotarse por plan/carrera, historial oficial o regla productiva previa."
)


# =============================================================================
# RESÚMENES Y DICTAMEN
# =============================================================================

total_filas = len(aud_anual)
filas_sin_codcli = len(p1_sin_codcli)
filas_con_codcli = total_filas - filas_sin_codcli

if "TIENE_EVIDENCIA_2025" in aud_anual.columns:
    filas_con_2025 = int(aud_anual["TIENE_EVIDENCIA_2025"].apply(to_bool).sum())
else:
    filas_con_2025 = 0

if "TIENE_EVIDENCIA_2026_MISMO_CODCLI" in aud_anual.columns:
    filas_con_2026 = int(aud_anual["TIENE_EVIDENCIA_2026_MISMO_CODCLI"].apply(to_bool).sum())
else:
    filas_con_2026 = 0

filas_dif_16_19 = len(p2_diferencias)
filas_ok_16_19 = total_filas - filas_dif_16_19
filas_estado_no_clasif = len(p3_filas_impactadas)
filas_acum_revisar = len(p4_acum_revisar)
filas_acum_ok = len(aud_acum) - filas_acum_revisar

if filas_sin_codcli > 0:
    dictamen_anual = "BLOQUEO_SIN_CODCLI"
elif filas_estado_no_clasif > 0:
    dictamen_anual = "BLOQUEO_ESTADOS_NO_CLASIFICADOS"
elif filas_dif_16_19 > 0:
    dictamen_anual = "REVISAR_DIFERENCIAS_COLUMNAS_16_19"
else:
    dictamen_anual = "OK_COLUMNAS_16_19_RESPALDADAS_2025"

if filas_acum_revisar > 0:
    dictamen_acumulado = "REVISAR_LOGICA_ACUMULADA"
else:
    dictamen_acumulado = "OK_COLUMNAS_20_21_RESPALDADAS"

dictamen_global = pd.DataFrame([
    {
        "PROCESO": PROCESO,
        "SUBPROYECTO": SUBPROYECTO,
        "ANIO_REFERENCIA": ANIO_REFERENCIA,
        "TOTAL_FILAS_5809": total_filas,
        "FILAS_CON_CODCLI": filas_con_codcli,
        "FILAS_SIN_CODCLI": filas_sin_codcli,
        "FILAS_CON_EVIDENCIA_2025": filas_con_2025,
        "FILAS_CON_EVIDENCIA_2026_MISMO_CODCLI": filas_con_2026,
        "FILAS_OK_COLUMNAS_16_19": filas_ok_16_19,
        "FILAS_DIFERENCIA_COLUMNAS_16_19": filas_dif_16_19,
        "FILAS_ESTADO_NO_CLASIFICADO": filas_estado_no_clasif,
        "FILAS_ACUMULADO_OK": filas_acum_ok,
        "FILAS_ACUMULADO_REVISAR": filas_acum_revisar,
        "DICTAMEN_ANUAL_2025": dictamen_anual,
        "DICTAMEN_ACUMULADO_2025": dictamen_acumulado,
        "DECLARACION_CARGA": "NO_LISTO_PARA_CARGA_DEFINITIVA",
    }
])

plan_resolucion = pd.DataFrame([
    {
        "ORDEN": 1,
        "PUNTO": "Filas sin CODCLI",
        "CANTIDAD": filas_sin_codcli,
        "HOJA_TRABAJO": "P1_SIN_CODCLI_CANDIDATOS",
        "ACCION": "Validar candidatos. Si no hay candidato, solicitar CODCLI a Docencia o base maestra.",
        "CRITERIO_CIERRE": "Cada fila debe quedar con CODCLI trazable o pendiente documentado.",
        "ESTADO": "REVISAR" if filas_sin_codcli > 0 else "OK",
    },
    {
        "ORDEN": 2,
        "PUNTO": "Diferencias columnas 16-19",
        "CANTIDAD": filas_dif_16_19,
        "HOJA_TRABAJO": "P2_DIFERENCIAS_16_19",
        "ACCION": "Revisar diferencias por semestre, unidades cursadas y aprobadas contra evidencia 2025.",
        "CRITERIO_CIERRE": "Cada diferencia debe quedar clasificada como error de mapeo, lógica, estado o pendiente.",
        "ESTADO": "REVISAR" if filas_dif_16_19 > 0 else "OK",
    },
    {
        "ORDEN": 3,
        "PUNTO": "Estados no clasificados",
        "CANTIDAD": filas_estado_no_clasif,
        "HOJA_TRABAJO": "P3_ESTADOS_NO_CLASIFICADOS",
        "ACCION": "Solicitar clasificación funcional APROBADO / NO_APROBADO. No clasificar automáticamente.",
        "CRITERIO_CIERRE": "Cada estado debe tener clasificación respaldada.",
        "ESTADO": "REVISAR" if filas_estado_no_clasif > 0 else "OK",
    },
    {
        "ORDEN": 4,
        "PUNTO": "Lógica acumulada columnas 20-21",
        "CANTIDAD": filas_acum_revisar,
        "HOJA_TRABAJO": "P4_ACUMULADO_REVISAR",
        "ACCION": "Revisar lógica acumulada separada del dictamen anual. Definir si se acota por plan/carrera/historial.",
        "CRITERIO_CIERRE": "Lógica acumulada documentada y revalidada.",
        "ESTADO": "REVISAR" if filas_acum_revisar > 0 else "OK",
    },
])

validaciones = pd.DataFrame([
    {
        "VALIDACION": "AUDITORIA_CORREGIDA_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(AUDITORIA_CORREGIDA),
    },
    {
        "VALIDACION": "PROMEDIOS_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(PROMEDIOS),
    },
    {
        "VALIDACION": "CSV_5809_EXISTE",
        "RESULTADO": "OK",
        "DETALLE": str(CSV_5809),
    },
    {
        "VALIDACION": "HOJAS_AUDITORIA_REQUERIDAS",
        "RESULTADO": "OK" if validacion_hojas_df["EXISTE_Y_TIENE_FILAS"].eq("SI").any() else "REVISAR",
        "DETALLE": "Ver hoja 01B_HOJAS_AUDITORIA",
    },
    {
        "VALIDACION": "FUENTES_ORIGINALES_MODIFICADAS",
        "RESULTADO": "OK",
        "DETALLE": "NO",
    },
    {
        "VALIDACION": "DECLARACION_CARGA",
        "RESULTADO": "REVISAR",
        "DETALLE": "NO_LISTO_PARA_CARGA_DEFINITIVA",
    },
])


# =============================================================================
# EXPORTAR EXCEL
# =============================================================================

excel_salida = RESULTADOS / "MAPEO_RESOLUCION_BLOQUEOS_5809_LOGICA_2025.xlsx"

with pd.ExcelWriter(excel_salida, engine="openpyxl") as writer:
    dictamen_global.to_excel(writer, sheet_name="00_DICTAMEN_GLOBAL", index=False)
    validaciones.to_excel(writer, sheet_name="01_VALIDACIONES", index=False)
    validacion_hojas_df.to_excel(writer, sheet_name="01B_HOJAS_AUDITORIA", index=False)
    plan_resolucion.to_excel(writer, sheet_name="02_PLAN_RESOLUCION", index=False)
    columnas_prev.to_excel(writer, sheet_name="03_COLUMNAS_DETECTADAS", index=False)
    p1_sin_codcli.to_excel(writer, sheet_name="P1_SIN_CODCLI_5809", index=False)
    p1_candidatos.to_excel(writer, sheet_name="P1_SIN_CODCLI_CANDIDATOS", index=False)
    p2_diferencias.to_excel(writer, sheet_name="P2_DIFERENCIAS_16_19", index=False)
    filas_revisar_anual.to_excel(writer, sheet_name="P2_FUENTE_FILAS_REVISAR", index=False)
    p3_estados.to_excel(writer, sheet_name="P3_ESTADOS_NO_CLASIFICADOS", index=False)
    p3_filas_impactadas.to_excel(writer, sheet_name="P3_FILAS_IMPACTADAS", index=False)
    p4_acum_revisar.to_excel(writer, sheet_name="P4_ACUMULADO_REVISAR", index=False)
    filas_revisar_acum.to_excel(writer, sheet_name="P4_FUENTE_FILAS_REVISAR", index=False)
    resumen_anual.to_excel(writer, sheet_name="RESUMEN_ANUAL_PREVIO", index=False)
    resumen_acum.to_excel(writer, sheet_name="RESUMEN_ACUMULADO_PREVIO", index=False)
    logica.to_excel(writer, sheet_name="LOGICA_AUDITADA", index=False)
    catalogo_estado.to_excel(writer, sheet_name="CATALOGO_ESTADO_COMPLETO", index=False)
    mapeo_codcli_rut.to_excel(writer, sheet_name="MAPEO_CODCLI_RUT_PREVIO", index=False)
    mapeo_promedios.to_excel(writer, sheet_name="MAPEO_PROMEDIOS_RECONSTRUIDO", index=False)

    pd.DataFrame([
        {
            "TIPO": "AUDITORIA_CORREGIDA",
            "RUTA": str(AUDITORIA_CORREGIDA),
            "SHA256": sha256(AUDITORIA_CORREGIDA),
        },
        {
            "TIPO": "PROMEDIOS",
            "RUTA": str(PROMEDIOS),
            "SHA256": sha256(PROMEDIOS),
        },
        {
            "TIPO": "CSV_5809",
            "RUTA": str(CSV_5809),
            "SHA256": sha256(CSV_5809),
        },
    ]).to_excel(writer, sheet_name="FUENTES", index=False)

aplicar_formato_excel(excel_salida)
hojas_generadas = validar_xlsx(excel_salida)


# =============================================================================
# INFORME / MANIFEST / COPIA ESCRITORIO
# =============================================================================

informe = REPORTES / "INFORME_MAPEO_RESOLUCION_BLOQUEOS_5809_LOGICA_2025.md"

informe.write_text(
    "# Mapeo resolución bloqueos 5809 lógica 2025\n\n"
    f"Proceso: {PROCESO}  \n"
    f"Subproyecto: {SUBPROYECTO}  \n"
    f"Año referencia: {ANIO_REFERENCIA}  \n\n"
    "## Dictamen\n\n"
    f"- Dictamen anual 2025: **{dictamen_anual}**\n"
    f"- Dictamen acumulado 2025: **{dictamen_acumulado}**\n"
    "- Declaración de carga: **NO_LISTO_PARA_CARGA_DEFINITIVA**\n\n"
    "## Indicadores principales\n\n"
    f"- Total filas 5809: {total_filas}\n"
    f"- Filas con CODCLI: {filas_con_codcli}\n"
    f"- Filas sin CODCLI: {filas_sin_codcli}\n"
    f"- Filas con evidencia 2025: {filas_con_2025}\n"
    f"- Filas con evidencia 2026 mismo CODCLI: {filas_con_2026}\n"
    f"- Filas OK columnas 16-19: {filas_ok_16_19}\n"
    f"- Filas con diferencia columnas 16-19: {filas_dif_16_19}\n"
    f"- Filas con estado no clasificado: {filas_estado_no_clasif}\n"
    f"- Filas acumulado revisar: {filas_acum_revisar}\n\n"
    "## Paso siguiente\n\n"
    "Trabajar las hojas en este orden:\n\n"
    "1. `P1_SIN_CODCLI_CANDIDATOS`\n"
    "2. `P2_DIFERENCIAS_16_19`\n"
    "3. `P3_ESTADOS_NO_CLASIFICADOS`\n"
    "4. `P4_ACUMULADO_REVISAR`\n\n"
    "Este archivo no corrige el 5809. Solo mapea cómo resolver los bloqueos.\n\n"
    f"- Excel: `{excel_salida}`\n"
    f"- Carpeta repo: `{SALIDA}`\n"
    f"- Carpeta Escritorio: `{CARPETA_ESCRITORIO}`\n",
    encoding="utf-8",
)

manifest = {
    "fecha": datetime.now().isoformat(),
    "proceso": PROCESO,
    "subproyecto": SUBPROYECTO,
    "anio_referencia": ANIO_REFERENCIA,
    "auditoria_corregida": str(AUDITORIA_CORREGIDA),
    "promedios": str(PROMEDIOS),
    "csv_5809": str(CSV_5809),
    "excel_salida": str(excel_salida),
    "informe": str(informe),
    "carpeta_repo": str(SALIDA),
    "carpeta_escritorio": str(CARPETA_ESCRITORIO),
    "total_filas_5809": total_filas,
    "filas_con_codcli": filas_con_codcli,
    "filas_sin_codcli": filas_sin_codcli,
    "filas_con_evidencia_2025": filas_con_2025,
    "filas_con_evidencia_2026_mismo_codcli": filas_con_2026,
    "filas_ok_columnas_16_19": filas_ok_16_19,
    "filas_diferencia_columnas_16_19": filas_dif_16_19,
    "filas_estado_no_clasificado": filas_estado_no_clasif,
    "filas_acumulado_revisar": filas_acum_revisar,
    "dictamen_anual_2025": dictamen_anual,
    "dictamen_acumulado_2025": dictamen_acumulado,
    "declaracion_carga": "NO_LISTO_PARA_CARGA_DEFINITIVA",
    "hojas_generadas": hojas_generadas,
    "fuentes_originales_modificadas": False,
}

manifest_path = MANIFEST_DIR / "manifest_mapeo_resolucion_bloqueos_5809_logica_2025.json"

manifest_path.write_text(
    dumps_json_safe(manifest),
    encoding="utf-8",
)

script_actual = Path(__file__).resolve()

for archivo in [excel_salida, informe, manifest_path, script_actual]:
    shutil.copy2(archivo, CARPETA_ESCRITORIO / archivo.name)

leeme = CARPETA_ESCRITORIO / "LEEME_MAPEO_RESOLUCION_BLOQUEOS_5809_2025.txt"

leeme.write_text(
    "MAPEO RESOLUCION BLOQUEOS 5809 LOGICA 2025 - AVANCE CURRICULAR SIES 2026\n"
    f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
    f"Dictamen anual 2025: {dictamen_anual}\n"
    f"Dictamen acumulado 2025: {dictamen_acumulado}\n"
    "Declaración de carga: NO_LISTO_PARA_CARGA_DEFINITIVA\n\n"
    "Hojas principales:\n"
    "- 02_PLAN_RESOLUCION\n"
    "- P1_SIN_CODCLI_CANDIDATOS\n"
    "- P2_DIFERENCIAS_16_19\n"
    "- P3_ESTADOS_NO_CLASIFICADOS\n"
    "- P4_ACUMULADO_REVISAR\n\n"
    "Advertencia:\n"
    "- Este archivo mapea cómo resolver bloqueos.\n"
    "- No modifica el 5809.\n"
    "- No modifica fuentes originales.\n",
    encoding="utf-8",
)

print()
print("=" * 120)
print("MAPEO RESOLUCION BLOQUEOS 5809 LOGICA 2025 — GENERADO")
print("=" * 120)
print("Fuentes originales modificadas: NO")
print("Declaración de carga: NO_LISTO_PARA_CARGA_DEFINITIVA")
print()
print("DICTAMEN GLOBAL")
print("-" * 120)
print(dictamen_global.to_string(index=False))
print()
print("PLAN DE RESOLUCION")
print("-" * 120)
print(plan_resolucion.to_string(index=False))
print()
print("VALIDACIONES")
print("-" * 120)
print(validaciones.to_string(index=False))
print()
print("ARCHIVOS GENERADOS")
print("-" * 120)
print(f"Excel:              {excel_salida}")
print(f"Informe:            {informe}")
print(f"Manifest:           {manifest_path}")
print(f"Script:             {script_actual}")
print(f"Carpeta repo:       {SALIDA}")
print(f"Carpeta escritorio: {CARPETA_ESCRITORIO}")
print("=" * 120)
