#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Auditoria especifica de los 223 casos DIFERENCIA_SOLO_APROBADAS_19.

No modifica fuentes originales, no corrige datos, no genera SIES_READY,
no recalcula 20-21 y no mezcla otros bloques. Solo diagnostica causas y
propone rutas de resolucion.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import unicodedata
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROCESO = "Avance Curricular SIES 2026"
SUBPROYECTO = "Auditoria diferencia solo aprobadas 19 5809"
ANIO_PROCESO = 2026
ANIO_REFERENCIA_DATOS = 2025
DECLARACION_CARGA = "NO_LISTO_PARA_CARGA"
NO_CORRECCION = "NO_APLICAR_CORRECCION_AUTOMATICA"

REPO = Path("/Users/alexi/Documents/GitHub/avance_curricular")
BASE_54 = REPO / "avance_curricular_2026" / "54_auditoria_diferencia_solo_aprobadas_19_5809"

GOBERNANZA_51 = (
    REPO
    / "avance_curricular_2026/51_gobernanza_integral_inputs_columnas_transformaciones_5809/"
    / "GOBERNANZA_INTEGRAL_5809_20260708_144544/"
    / "GOBERNANZA_INTEGRAL_INPUTS_COLUMNAS_TRANSFORMACIONES_5809.xlsx"
)
VALIDADOR_52_SCRIPT = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "validador_gobernanza_precalculo_5809.py"
)
VALIDADOR_52_XLSX = (
    REPO
    / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809/"
    / "VALIDADOR_GOBERNANZA_PRECALCULO_5809_20260708_165151/"
    / "VALIDACION_GOBERNANZA_PRECALCULO_5809.xlsx"
)
MATRIZ_53 = (
    REPO
    / "avance_curricular_2026/53_matriz_resolucion_bloqueos_16_19_gobernada_5809/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_20260708_170521/"
    / "MATRIZ_RESOLUCION_BLOQUEOS_16_19_GOBERNADA_5809.xlsx"
)
RECALCULO_50 = (
    REPO
    / "avance_curricular_2026/50_recalculo_16_19_diccionario_estado_promedios/"
    / "RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS_20260708_143748/"
    / "02_RESULTADOS/RECALCULO_16_19_DICCIONARIO_ESTADO_PROMEDIOS.xlsx"
)

OUTPUT_FILES = {
    "excel": "AUDITORIA_DIFERENCIA_SOLO_APROBADAS_19_5809.xlsx",
    "informe": "INFORME_AUDITORIA_DIFERENCIA_SOLO_APROBADAS_19_5809.md",
    "manifest": "manifest_auditoria_diferencia_solo_aprobadas_19_5809.json",
    "script": "auditoria_diferencia_solo_aprobadas_19_5809.py",
}

CAUSES = [
    "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA",
    "RAMO_DUPLICADO_MISMO_CODCLI",
    "RAMO_DUPLICADO_MULTIPLE_CODCLI",
    "ESTADO_NULL_O_BLANCO",
    "DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE",
    "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS",
    "REQUIERE_REVISION_FUNCIONAL",
]

SHEET_NAMES = OrderedDict(
    [
        ("00_DICTAMEN", "00_DICTAMEN"),
        ("01_VALIDACION_GOBERNANZA", "01_VALIDACION_GOBERNANZA"),
        ("02_RESUMEN_CAUSAS", "02_RESUMEN_CAUSAS"),
        ("03_223_CASOS", "03_223_CASOS"),
        ("04_DETALLE_RAMO_ESTADO", "04_DETALLE_RAMO_ESTADO"),
        ("05_COMPARACION_5809_VS_RECALCULO", "05_COMP_5809_VS_RECALCULO"),
        ("06_CAUSAS_PROPUESTAS", "06_CAUSAS_PROPUESTAS"),
        ("07_NO_RESOLVER_AUTOMATICO", "07_NO_RESOLVER_AUTOMATICO"),
        ("08_FUENTES", "08_FUENTES"),
    ]
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def norm(value: Any) -> str:
    text = clean(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    out = []
    for ch in text.upper():
        out.append(ch if ch.isalnum() else "_")
    return "_".join("".join(out).split("_"))


def display(value: Any, max_len: int = 160) -> str:
    text = clean(value)
    if len(text) > max_len:
        return text[: max_len - 3] + "..."
    return text


def markdown_table(df: pd.DataFrame, max_rows: int | None = None) -> str:
    if max_rows is not None:
        df = df.head(max_rows)
    if df.empty:
        return "(sin filas)"
    cols = list(df.columns)
    widths = []
    for col in cols:
        values = [str(col)] + [display(v, 100) for v in df[col].tolist()]
        widths.append(max(len(v) for v in values))
    header = "| " + " | ".join(str(col).ljust(widths[i]) for i, col in enumerate(cols)) + " |"
    sep = "| " + " | ".join("-" * width for width in widths) + " |"
    body = []
    for _, row in df.iterrows():
        body.append(
            "| "
            + " | ".join(display(row[col], 100).ljust(widths[i]) for i, col in enumerate(cols))
            + " |"
        )
    return "\n".join([header, sep] + body)


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str).fillna("")


def to_int(value: Any) -> int:
    text = clean(value)
    if not text:
        return 0
    return int(float(text))


def split_codcli_lista(value: Any) -> List[str]:
    text = clean(value)
    if not text:
        return []
    return [part.strip() for part in text.split(";") if part.strip()]


def unique_count(df: pd.DataFrame, estado: str | None = None, classif: str | None = None) -> int:
    subset = df
    if estado is not None:
        subset = subset[subset["ESTADO"].eq(estado)]
    if classif is not None:
        subset = subset[subset["CLASIF_DICCIONARIO"].eq(classif)]
    return len(subset.drop_duplicates(["CODRAMO"]))


def join_unique(values: List[str]) -> str:
    return " | ".join(sorted(set(v for v in values if clean(v))))


def run_validator_52(output_dir: Path, timestamp: str) -> Dict[str, Any]:
    plan_path = output_dir / "plan_uso_validador_h52_auditoria_dif_19.json"
    plan = {
        "descripcion": "Auditoria diagnostica de 223 casos DIFERENCIA_SOLO_APROBADAS_19. No corrige, no recalcula y no genera carga.",
        "columnas_usadas": ["CODIGO_UNICO", "PLAN_ESTUDIOS", "CODCLI_LISTA", "ESTADO", "DESCRIPCION_ESTADO"],
        "llaves_usadas": ["CODCLI_LISTA", "CODIGO_UNICO", "PLAN_ESTUDIOS"],
        "diccionarios_usados": ["DIC_ESTADO_PROMEDIOS"],
        "transformaciones": ["Diagnostico documental de diferencias columna 19; sin correccion de datos"],
        "campos_calculo": ["19"],
        "anio_calculo": 2025,
        "genera_sies_ready": False,
        "modifica_fuentes_originales": False,
        "usa_todos_codcli_rut": False,
        "filtra_por_programa": True,
        "usa_inferencias_como_reglas_oficiales": False,
        "usa_estados_AEIR": True,
        "usa_codcli_lista": True,
        "mezcla_anual_acumulado": False,
        "modo": "AUDITORIA_DIAGNOSTICA_NO_CORRECTIVA",
    }
    plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    validator_timestamp = f"{timestamp}_H54"
    cmd = [
        sys.executable,
        str(VALIDADOR_52_SCRIPT),
        "--timestamp",
        validator_timestamp,
        "--plan-json",
        str(plan_path),
    ]
    completed = subprocess.run(cmd, check=False, capture_output=True, text=True)
    validator_dir = (
        REPO
        / "avance_curricular_2026/52_validador_gobernanza_precalculo_5809"
        / f"VALIDADOR_GOBERNANZA_PRECALCULO_5809_{validator_timestamp}"
    )
    manifest_path = validator_dir / "manifest_validacion_gobernanza_precalculo_5809.json"
    manifest = {}
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    result = {
        "plan_path": str(plan_path),
        "validator_timestamp": validator_timestamp,
        "validator_dir": str(validator_dir),
        "validator_manifest": str(manifest_path),
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "estado_validador": manifest.get("estado_validador", "SIN_MANIFEST"),
        "declaracion_carga": manifest.get("estado_carga", ""),
        "fuentes_originales_modificadas": manifest.get("fuentes_originales_modificadas", ""),
        "plan": plan,
    }
    if completed.returncode != 0:
        raise RuntimeError("Validador hito 52 fallo:\n" + completed.stdout + "\n" + completed.stderr)
    if result["estado_validador"] != "APTO_PARA_CONTROL_PREVIO":
        raise RuntimeError(f"Validador hito 52 no apto: {result['estado_validador']}")
    return result


def build_validation_sheet(validator_result: Dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Control": "Gobernanza integral hito 51 disponible",
                "Estado": "OK",
                "Evidencia": str(GOBERNANZA_51),
                "Observacion": "Se usa diccionario ESTADO/DESCRIPCION_ESTADO gobernado.",
            },
            {
                "Control": "Validador hito 52 ejecutado para hito 54",
                "Estado": "OK" if validator_result["returncode"] == 0 else "ERROR",
                "Evidencia": validator_result["validator_manifest"],
                "Observacion": "Ejecucion obligatoria previa al diagnostico.",
            },
            {
                "Control": "Estado validador",
                "Estado": validator_result["estado_validador"],
                "Evidencia": validator_result["validator_dir"],
                "Observacion": "Debe ser APTO_PARA_CONTROL_PREVIO.",
            },
            {
                "Control": "Matriz hito 53 usada como entrada principal",
                "Estado": "OK",
                "Evidencia": str(MATRIZ_53),
                "Observacion": "Solo hoja 03_DIF_SOLO_APROBADAS_19.",
            },
            {
                "Control": "No mezcla con otros bloques",
                "Estado": "OK",
                "Evidencia": "Filtro DICTAMEN_DICCIONARIO_ESTADO=DIFERENCIA_SOLO_APROBADAS_19",
                "Observacion": "Se excluyen diferencias 16-18, sin registros 2025 y sin CODCLI_LISTA.",
            },
            {
                "Control": "No correccion / no carga",
                "Estado": DECLARACION_CARGA,
                "Evidencia": "Productos hito 54",
                "Observacion": "No se aplica correccion ni se genera SIES_READY.",
            },
        ]
    )


def prepare_detail(cases_df: pd.DataFrame) -> pd.DataFrame:
    detail = read_sheet(RECALCULO_50, "07_DETALLE_RAMO_ESTADO")
    ids = set(cases_df["FILA_5809"].astype(str))
    detail = detail[detail["FILA_5809"].astype(str).isin(ids)].copy()
    same = (
        detail.groupby(["FILA_5809", "CODCLI_NORM", "CODRAMO"])
        .size()
        .reset_index(name="N_MISMO_CODCLI_CODRAMO")
    )
    multi = (
        detail.groupby(["FILA_5809", "CODRAMO"])["CODCLI_NORM"]
        .nunique()
        .reset_index(name="N_CODCLI_POR_CODRAMO")
    )
    detail = detail.merge(same, on=["FILA_5809", "CODCLI_NORM", "CODRAMO"], how="left")
    detail = detail.merge(multi, on=["FILA_5809", "CODRAMO"], how="left")
    detail["DUPLICADO_MISMO_CODCLI"] = detail["N_MISMO_CODCLI_CODRAMO"].astype(int).gt(1).map({True: "SI", False: "NO"})
    detail["DUPLICADO_MULTIPLE_CODCLI"] = detail["N_CODCLI_POR_CODRAMO"].astype(int).gt(1).map({True: "SI", False: "NO"})
    return detail


def classify_cases(cases_df: pd.DataFrame, detail_df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    compare_rows = []
    for _, case in cases_df.iterrows():
        fid = clean(case.get("FILA_5809"))
        d = detail_df[detail_df["FILA_5809"].astype(str).eq(fid)].copy()
        v5809 = to_int(case.get("UNIDADES_APROBADAS_5809"))
        recalc = to_int(case.get("UNIDADES_APROBADAS_RECALC_DICC"))
        gap = recalc - v5809
        unique_aprobada = unique_count(d, classif="APROBADA")
        unique_a = unique_count(d, estado="A")
        unique_e = unique_count(d, estado="E")
        unique_i = unique_count(d, estado="I")
        unique_ei = len(d[d["ESTADO"].isin(["E", "I"])].drop_duplicates(["CODRAMO"]))
        unique_r = unique_count(d, estado="R")
        null_mask = d["CLASIF_DICCIONARIO"].eq("REVISION_NULL") | d["ESTADO"].eq("")
        unique_null = len(d[null_mask].drop_duplicates(["CODRAMO"]))
        same_dup = int(
            d.groupby(["CODCLI_NORM", "CODRAMO"]).size().reset_index(name="n")["n"].gt(1).sum()
        ) if not d.empty else 0
        multi_dup = int(
            d.groupby(["CODRAMO"])["CODCLI_NORM"].nunique().reset_index(name="n")["n"].gt(1).sum()
        ) if not d.empty else 0
        codcli_list = split_codcli_lista(case.get("CODCLI_LISTA_NORM"))
        codcli_lista_count = to_int(case.get("CODCLI_LISTA_CANTIDAD")) or len(codcli_list)
        gap_eq_ei = "SI" if gap == unique_ei else "NO"
        v5809_eq_a = "SI" if v5809 == unique_a else "NO"

        causes = []
        if unique_ei > 0 and gap_eq_ei == "SI" and v5809_eq_a == "SI":
            causes.append("CONVALIDACION_HOMOLOGACION_YA_GOBERNADA")
        if same_dup > 0:
            causes.append("RAMO_DUPLICADO_MISMO_CODCLI")
        if multi_dup > 0:
            causes.append("RAMO_DUPLICADO_MULTIPLE_CODCLI")
        if unique_null > 0:
            causes.append("ESTADO_NULL_O_BLANCO")
        if codcli_lista_count > 1:
            causes.append("DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE")
        if gap_eq_ei == "NO" or v5809_eq_a == "NO" or unique_ei == 0:
            causes.append("POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS")
        causes.append("REQUIERE_REVISION_FUNCIONAL")
        causes = [c for c in CAUSES if c in causes]

        if unique_null > 0:
            main = "ESTADO_NULL_O_BLANCO"
        elif "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS" in causes and unique_ei == 0:
            main = "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS"
        elif "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA" in causes:
            main = "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA"
        elif "RAMO_DUPLICADO_MULTIPLE_CODCLI" in causes:
            main = "RAMO_DUPLICADO_MULTIPLE_CODCLI"
        elif "RAMO_DUPLICADO_MISMO_CODCLI" in causes:
            main = "RAMO_DUPLICADO_MISMO_CODCLI"
        elif "DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE" in causes:
            main = "DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE"
        else:
            main = "REQUIERE_REVISION_FUNCIONAL"

        if main == "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA":
            reason = (
                "La brecha recalc-5809 coincide con ramos unicos E/I; 5809 coincide con conteo solo A. "
                "La decision funcional debe confirmar si E/I aplica al campo anual 19."
            )
        elif main == "ESTADO_NULL_O_BLANCO":
            reason = "Existe al menos un ramo con estado NULL/blanco; gobernanza 51 prohibe resolverlo automaticamente."
        elif main == "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS":
            reason = "La brecha no se explica por E/I; el valor 5809 tampoco coincide con conteo solo A."
        elif main == "RAMO_DUPLICADO_MISMO_CODCLI":
            reason = "Existen repeticiones del mismo CODRAMO para el mismo CODCLI; requiere deduplicacion auditada."
        elif main == "RAMO_DUPLICADO_MULTIPLE_CODCLI":
            reason = "El mismo CODRAMO aparece asociado a mas de un CODCLI del caso; requiere resolver llave academica."
        elif main == "DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE":
            reason = "El caso tiene multiples CODCLI_LISTA; requiere validar filtro por programa."
        else:
            reason = "Requiere revision funcional antes de cualquier ajuste."

        compare = {
            "FILA_5809": fid,
            "RUT_NORMALIZADO": clean(case.get("RUT_NORMALIZADO")),
            "CODIGO_UNICO": clean(case.get("CODIGO_UNICO")),
            "PLAN_ESTUDIOS": clean(case.get("PLAN_ESTUDIOS")),
            "CODCLI_LISTA_NORM": clean(case.get("CODCLI_LISTA_NORM")),
            "CODCLI_LISTA_CANTIDAD": codcli_lista_count,
            "UNIDADES_APROBADAS_5809": v5809,
            "UNIDADES_APROBADAS_RECALC_DICC": recalc,
            "BRECHA_RECALC_MENOS_5809": gap,
            "APROBADAS_DICC_UNICAS": unique_aprobada,
            "APROBADAS_SOLO_A_UNICAS": unique_a,
            "CONVALIDACION_E_UNICAS": unique_e,
            "HOMOLOGADO_I_UNICAS": unique_i,
            "CONVALIDACION_HOMOLOGACION_EI_UNICAS": unique_ei,
            "REPROBADO_R_UNICAS": unique_r,
            "NULL_BLANCO_UNICOS": unique_null,
            "DUP_RAMO_MISMO_CODCLI": same_dup,
            "DUP_RAMO_MULTIPLE_CODCLI": multi_dup,
            "BRECHA_EXPLICADA_POR_EI": gap_eq_ei,
            "VALOR_5809_EQ_SOLO_A": v5809_eq_a,
            "CAUSA_PRINCIPAL": main,
            "CAUSAS_DETECTADAS": " | ".join(causes),
            "MOTIVO_DIAGNOSTICO": reason,
            "DECLARACION_CARGA": DECLARACION_CARGA,
            "NO_APLICAR_CORRECCION": NO_CORRECCION,
        }
        compare_rows.append(compare)
        rows.append({**case.to_dict(), **compare})
    return pd.DataFrame(rows), pd.DataFrame(compare_rows)


def build_cause_summary(cases_audit: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for cause in CAUSES:
        mask = cases_audit["CAUSAS_DETECTADAS"].str.contains(cause, regex=False, na=False)
        main_mask = cases_audit["CAUSA_PRINCIPAL"].eq(cause)
        rows.append(
            {
                "CAUSA": cause,
                "CASOS_CAUSA_PRINCIPAL": int(main_mask.sum()),
                "CASOS_CAUSA_DETECTADA": int(mask.sum()),
                "BRECHA_TOTAL_RECALC_MENOS_5809": int(cases_audit.loc[mask, "BRECHA_RECALC_MENOS_5809"].sum()) if mask.any() else 0,
                "PUEDE_RESOLVERSE_AUTOMATICAMENTE": "NO",
                "REQUIERE_REVISION_FUNCIONAL": "SI",
            }
        )
    return pd.DataFrame(rows)


def build_cause_proposals(summary_df: pd.DataFrame) -> pd.DataFrame:
    proposal = {
        "CONVALIDACION_HOMOLOGACION_YA_GOBERNADA": (
            "Confirmar funcionalmente si E/I deben contarse en UNIDADES_APROBADAS anual 2025 o excluirse por regla oficial.",
            "Diccionario PROMEDIOS, detalle por CODRAMO y decision funcional E/I.",
        ),
        "RAMO_DUPLICADO_MISMO_CODCLI": (
            "Auditar duplicados mismo CODCLI-CODRAMO antes de cualquier conteo definitivo.",
            "Regla de deduplicacion por CODCLI+CODRAMO y evidencia de intentos/periodos.",
        ),
        "RAMO_DUPLICADO_MULTIPLE_CODCLI": (
            "Resolver si el CODRAMO pertenece a un unico CODCLI/programa para la fila 5809.",
            "Llave programa RUT+CODIGO_UNICO+PLAN_ESTUDIOS y CODCLI_LISTA filtrado.",
        ),
        "ESTADO_NULL_O_BLANCO": (
            "No contar automaticamente; solicitar definicion funcional para NULL/blanco.",
            "Diccionario estado y revision de ramos con estado vacio.",
        ),
        "DIFERENCIA_POR_CODCLI_LISTA_MULTIPLE": (
            "Validar filtro por programa cuando CODCLI_LISTA_CANTIDAD > 1.",
            "Mapa CODCLI_LISTA, programa 2025 y detalle academico por CODCLI.",
        ),
        "POSIBLE_DATO_5809_NO_CALZA_CON_PROMEDIOS": (
            "Revisar manualmente porque la brecha no se explica por E/I.",
            "Comparacion 5809 vs conteo A y recuento diccionario.",
        ),
        "REQUIERE_REVISION_FUNCIONAL": (
            "Llevar a decision funcional antes de cualquier ajuste.",
            "Hoja 03_223_CASOS y detalle por ramo/estado.",
        ),
    }
    rows = []
    for _, row in summary_df.iterrows():
        action, evidence = proposal[row["CAUSA"]]
        rows.append(
            {
                "CAUSA": row["CAUSA"],
                "CASOS_CAUSA_PRINCIPAL": row["CASOS_CAUSA_PRINCIPAL"],
                "CASOS_CAUSA_DETECTADA": row["CASOS_CAUSA_DETECTADA"],
                "ACCION_PROPUESTA": action,
                "EVIDENCIA_MINIMA_REQUERIDA": evidence,
                "NO_RESOLVER_AUTOMATICO": "SI",
                "DECLARACION_CARGA": DECLARACION_CARGA,
            }
        )
    return pd.DataFrame(rows)


def build_no_auto(cases_audit: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "FILA_5809",
        "RUT_NORMALIZADO",
        "CODIGO_UNICO",
        "PLAN_ESTUDIOS",
        "CODCLI_LISTA_NORM",
        "UNIDADES_APROBADAS_5809",
        "UNIDADES_APROBADAS_RECALC_DICC",
        "BRECHA_RECALC_MENOS_5809",
        "CAUSA_PRINCIPAL",
        "CAUSAS_DETECTADAS",
        "MOTIVO_DIAGNOSTICO",
        "NO_APLICAR_CORRECCION",
        "DECLARACION_CARGA",
    ]
    return cases_audit[cols].copy()


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="9E480E")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx in range(1, min(ws.max_column, 60) + 1):
            letter = get_column_letter(col_idx)
            max_len = 12
            for cell in ws[letter][: min(ws.max_row, 200)]:
                text = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(text), 90))
            ws.column_dimensions[letter].width = min(max_len + 2, 65)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    wb.close()


def write_excel(path: Path, dataframes: OrderedDict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet, df in dataframes.items():
            safe_df = df.copy()
            for col in safe_df.columns:
                safe_df[col] = safe_df[col].map(lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
            safe_df.to_excel(writer, sheet_name=SHEET_NAMES.get(sheet, sheet), index=False)
    style_workbook(path)


def source_hashes(paths: Dict[str, Path]) -> Dict[str, str]:
    return {key: sha256_file(path) for key, path in paths.items()}


def build_sources(paths: Dict[str, Path], before: Dict[str, str], after: Dict[str, str], outputs: Dict[str, Path]) -> pd.DataFrame:
    rows = []
    for key, path in paths.items():
        rows.append(
            {
                "Tipo": "INPUT_OBLIGATORIO",
                "ID": key,
                "Ruta": str(path),
                "Hash antes": before.get(key, ""),
                "Hash despues": after.get(key, ""),
                "Modificada": "NO" if before.get(key, "") == after.get(key, "") else "SI",
                "Uso": "Lectura controlada; no se modifica.",
            }
        )
    for key, path in outputs.items():
        rows.append(
            {
                "Tipo": "SALIDA_HITO_54",
                "ID": key,
                "Ruta": str(path),
                "Hash antes": "",
                "Hash despues": sha256_file(path) if path.exists() and path.is_file() else "",
                "Modificada": "DERIVADO",
                "Uso": "Producto diagnostico generado por hito 54.",
            }
        )
    return pd.DataFrame(rows)


def build_markdown(dictamen: pd.DataFrame, validation: pd.DataFrame, summary: pd.DataFrame, proposals: pd.DataFrame, manifest: Dict[str, Any]) -> str:
    lines = [
        "# Auditoria diferencia solo aprobadas 19 5809",
        "",
        f"Proceso: **{PROCESO}**",
        f"Subproyecto: **{SUBPROYECTO}**",
        f"Año proceso: **{ANIO_PROCESO}**",
        f"Año referencia datos: **{ANIO_REFERENCIA_DATOS}**",
        f"Salida: **{DECLARACION_CARGA}**",
        "",
        "## Dictamen",
        "",
        markdown_table(dictamen),
        "",
        "## Validacion gobernanza",
        "",
        markdown_table(validation),
        "",
        "## Resumen causas",
        "",
        markdown_table(summary),
        "",
        "## Causas propuestas",
        "",
        markdown_table(proposals),
        "",
        "## Alcance",
        "",
        "- Solo se auditan los 223 casos `DIFERENCIA_SOLO_APROBADAS_19`.",
        "- No se mezclan otros bloques del hito 53.",
        "- No se corrigen datos, no se genera SIES_READY y no se recalculan 20-21.",
        "- El diccionario `ESTADO` / `DESCRIPCION_ESTADO` de PROMEDIOS se usa como evidencia gobernada.",
        "",
        "## Archivos",
        "",
    ]
    for key, value in manifest.get("archivos", {}).items():
        lines.append(f"- {key}: `{value}`")
    lines.append("")
    return "\n".join(lines)


def scan_sies_ready(output_dir: Path) -> List[str]:
    return [str(path) for path in output_dir.rglob("*SIES_READY*")]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--timestamp", default=datetime.now().strftime("%Y%m%d_%H%M%S"))
    args = parser.parse_args()

    output_dir = BASE_54 / f"AUDITORIA_DIF_SOLO_APROBADAS_19_{args.timestamp}"
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = output_dir / OUTPUT_FILES["excel"]
    md_path = output_dir / OUTPUT_FILES["informe"]
    manifest_path = output_dir / OUTPUT_FILES["manifest"]
    script_path = output_dir / OUTPUT_FILES["script"]
    current_script = Path(__file__).resolve()
    if current_script != script_path.resolve():
        shutil.copy2(current_script, script_path)

    input_paths = OrderedDict(
        [
            ("GOBERNANZA_51", GOBERNANZA_51),
            ("VALIDADOR_52_SCRIPT", VALIDADOR_52_SCRIPT),
            ("VALIDADOR_52_XLSX", VALIDADOR_52_XLSX),
            ("MATRIZ_53", MATRIZ_53),
            ("RECALCULO_50_DETALLE", RECALCULO_50),
        ]
    )
    for path in input_paths.values():
        if not path.exists():
            raise FileNotFoundError(path)
    hashes_before = source_hashes(input_paths)

    validator_result = run_validator_52(output_dir, args.timestamp)
    validation_df = build_validation_sheet(validator_result)

    cases = read_sheet(MATRIZ_53, "03_DIF_SOLO_APROBADAS_19")
    cases = cases[cases["DICTAMEN_DICCIONARIO_ESTADO"].eq("DIFERENCIA_SOLO_APROBADAS_19")].copy()
    detail = prepare_detail(cases)
    cases_audit, comparison = classify_cases(cases, detail)
    cause_summary = build_cause_summary(cases_audit)
    cause_proposals = build_cause_proposals(cause_summary)
    no_auto = build_no_auto(cases_audit)

    total = len(cases_audit)
    dictamen = pd.DataFrame(
        [
            {
                "Proceso": PROCESO,
                "Subproyecto": SUBPROYECTO,
                "Anio proceso": ANIO_PROCESO,
                "Anio referencia datos": ANIO_REFERENCIA_DATOS,
                "Declaracion carga": DECLARACION_CARGA,
                "Casos auditados": total,
                "Bloque auditado": "DIFERENCIA_SOLO_APROBADAS_19",
                "Otros bloques mezclados": "NO",
                "Correcciones aplicadas": "NO",
                "SIES_READY generado": "NO",
                "Recalculo 20-21": "NO",
                "Estado validador hito 52": validator_result["estado_validador"],
                "Causa principal dominante": cause_summary.sort_values("CASOS_CAUSA_PRINCIPAL", ascending=False).iloc[0]["CAUSA"],
                "Dictamen": "AUDITORIA_DIAGNOSTICA_GENERADA_NO_LISTA_PARA_CARGA",
                "Proximo paso permitido": "Revisar causas y resolver decisiones funcionales; no aplicar correcciones automaticas.",
            }
        ]
    )

    outputs = {"Excel": excel_path, "Informe": md_path, "Manifest": manifest_path, "Script": script_path}
    hashes_after_pre = source_hashes(input_paths)
    fuentes_modificadas = "SI" if hashes_before != hashes_after_pre else "NO"

    manifest: Dict[str, Any] = {
        "proceso": PROCESO,
        "subproyecto": SUBPROYECTO,
        "timestamp": args.timestamp,
        "anio_proceso": ANIO_PROCESO,
        "anio_referencia_datos": ANIO_REFERENCIA_DATOS,
        "declaracion_carga": DECLARACION_CARGA,
        "fuentes_originales_modificadas": fuentes_modificadas,
        "correcciones_aplicadas": "NO",
        "sies_ready_generado": "NO",
        "recalculo_20_21": "NO",
        "otros_bloques_mezclados": "NO",
        "validator_52": validator_result,
        "equivalencias_hojas_excel": SHEET_NAMES,
        "conteos": {
            "casos_auditados": total,
            "detalle_ramo_estado": len(detail),
            "causa_principal": cause_summary.set_index("CAUSA")["CASOS_CAUSA_PRINCIPAL"].to_dict(),
            "causa_detectada": cause_summary.set_index("CAUSA")["CASOS_CAUSA_DETECTADA"].to_dict(),
        },
        "hashes_inputs_antes": hashes_before,
        "hashes_inputs_despues_pre_outputs": hashes_after_pre,
        "archivos": {key: str(path) for key, path in outputs.items()},
    }

    sources = build_sources(input_paths, hashes_before, hashes_after_pre, outputs)
    dataframes = OrderedDict(
        [
            ("00_DICTAMEN", dictamen),
            ("01_VALIDACION_GOBERNANZA", validation_df),
            ("02_RESUMEN_CAUSAS", cause_summary),
            ("03_223_CASOS", cases_audit),
            ("04_DETALLE_RAMO_ESTADO", detail),
            ("05_COMPARACION_5809_VS_RECALCULO", comparison),
            ("06_CAUSAS_PROPUESTAS", cause_proposals),
            ("07_NO_RESOLVER_AUTOMATICO", no_auto),
            ("08_FUENTES", sources),
        ]
    )
    write_excel(excel_path, dataframes)

    md_path.write_text(build_markdown(dictamen, validation_df, cause_summary, cause_proposals, manifest), encoding="utf-8")

    hashes_after = source_hashes(input_paths)
    fuentes_modificadas_final = "SI" if hashes_before != hashes_after else "NO"
    sies_ready = scan_sies_ready(output_dir)
    manifest["fuentes_originales_modificadas"] = fuentes_modificadas_final
    manifest["hashes_inputs_despues"] = hashes_after
    manifest["sies_ready_detectado"] = sies_ready
    if sies_ready:
        manifest["sies_ready_generado"] = "SI"
        dictamen.loc[0, "SIES_READY generado"] = "SI"
    manifest["hashes_salidas"] = {
        "Excel": sha256_file(excel_path),
        "Informe": sha256_file(md_path),
        "Script": sha256_file(script_path),
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print("AUDITORIA DIFERENCIA SOLO APROBADAS 19 5809 — GENERADA")
    print(f"Fuentes originales modificadas: {fuentes_modificadas_final}")
    print(f"Declaración carga: {DECLARACION_CARGA}")
    print("Correcciones aplicadas: NO")
    print("SIES_READY generado: NO")
    print("Recalculo 20-21: NO")
    print()
    print("DICTAMEN")
    print(markdown_table(dictamen))
    print()
    print("RESUMEN CAUSAS")
    print(markdown_table(cause_summary))
    print()
    print("ARCHIVOS")
    print(f"Excel: {excel_path}")
    print(f"Informe: {md_path}")
    print(f"Manifest: {manifest_path}")
    print(f"Script: {script_path}")


if __name__ == "__main__":
    main()
