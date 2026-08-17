#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

PRECARGA = ROOT / "02_precarga_pes" / "reporte_dinamico_5913_2026-08-12_08:22:23.csv"
ENTREGA_DIR = ROOT / "03_fuentes_institucionales" / "entregas_docencia" / "20260814_entrega_02"
DOCENCIA_MODIFICABLES = Path(os.environ.get(
    "DOCENCIA_MODIFICABLES",
    ENTREGA_DIR / "reporte_dinamico_5913_2026-08-12_14_52_50 Modificables_lm.xlsx",
))
DOCENCIA_VACANTES = Path(os.environ.get(
    "DOCENCIA_VACANTES",
    ENTREGA_DIR / "Vacantes 1° y 2° Sem_carreras_vigentes.xlsx",
))
DOCENCIA_HOJA = os.environ.get("DOCENCIA_HOJA", "")

OUT_TRABAJO = ROOT / "05_datos_trabajo"
OUT_VALIDACIONES = ROOT / "06_validaciones"
OUT_PENDIENTES = ROOT / "12_pendientes"
OUT_LOGS = ROOT / "08_logs"

NO_CARGA = {"CODIGO_IES_NUM", "CANTIDAD_MATRICULA_DFE", "CANTIDAD_BENEFICIO_DFE"}

CAMPOS_MODIFICABLES = [
    "REGIMEN",
    "DURACION_REGIMEN",
    "ACREDITACION",
    "REQUISITO_INGRESO",
    "SEMESTRES_RECONOCIDOS",
    "AREA_ADMIN_DERECHO",
    "AREA_AGRI_SILVI_PESCA_VET",
    "AREA_ARTES_HUMANIDADES",
    "AREA_CIENCIAS_NAT_MAT_ESTAD",
    "AREA_CS_SOCIAL_PERIODISMO_INFO",
    "AREA_EDUCACION",
    "AREA_INGE_INDUSTRIA_CONSTRUC",
    "AREA_SALUD_BIENESTAR",
    "AREA_SERVICIOS",
    "AREA_TECNO_INFO_COMUNICA",
    "VACANTES_PRIMER_SEMESTRE",
    "VACANTES_SEGUNDO_SEMESTRE",
    "FECHA_ADMISION_INICIAL",
    "ENLACE_INFO_PROGRAMA",
    "LICENCIA_ENS_MEDIA",
    "NOTAS_ENS_MEDIA",
    "PROMEDIO_MIN_ENS_MEDIA",
    "RECONOCIMIENTOS_APREN_PREVIOS",
    "EXPERIENCIA_LABORAL",
    "MAIL_DIFUSION_CARRERA",
    "FORMATO_VALOR",
    "VALOR_MATRICULA_ANUAL",
    "COSTO_TITULACION",
    "VALOR_CERTIFICADO_DIPLOMA",
    "ARANCEL_ANUAL",
    "VIGENCIA_CARRERA",
]

OBLIGATORIOS_VIGENCIA_1 = [
    "VACANTES_PRIMER_SEMESTRE",
    "VACANTES_SEGUNDO_SEMESTRE",
    "FECHA_ADMISION_INICIAL",
    "ENLACE_INFO_PROGRAMA",
    "LICENCIA_ENS_MEDIA",
    "NOTAS_ENS_MEDIA",
    "PROMEDIO_MIN_ENS_MEDIA",
    "RECONOCIMIENTOS_APREN_PREVIOS",
    "EXPERIENCIA_LABORAL",
    "MAIL_DIFUSION_CARRERA",
    "FORMATO_VALOR",
    "VALOR_MATRICULA_ANUAL",
    "COSTO_TITULACION",
    "VALOR_CERTIFICADO_DIPLOMA",
    "ARANCEL_ANUAL",
]

LLAVE_ORDEN = [
    "CODIGO_IES_NUM",
    "COD_SEDE",
    "NOMBRE_SEDE",
    "COD_CARRERA",
    "NOMBRE_CARRERA",
    "VERSION",
]

MAPAS = {
    "REGIMEN": {
        "SEMESTRES": "1",
        "TRIMESTRES": "2",
    },
    "ACREDITACION": {
        "SI": "1",
        "SÍ": "1",
        "ACREDITADA": "1",
        "NO": "2",
        "NO ACREDITADA": "2",
    },
    "REQUISITO_INGRESO": {
        "EDUCACION MEDIA": "1",
        "EDUCACIÓN MEDIA": "1",
        "TECNICO DE NIVEL SUPERIOR": "2",
        "TÉCNICO DE NIVEL SUPERIOR": "2",
        "BACHILLERATO": "3",
        "LICENCIATURA": "4",
    },
    "LICENCIA_ENS_MEDIA": {
        "SI": "1",
        "SÍ": "1",
        "NO": "2",
    },
    "NOTAS_ENS_MEDIA": {
        "SI": "1",
        "SÍ": "1",
        "NO": "2",
    },
    "RECONOCIMIENTOS_APREN_PREVIOS": {
        "SI": "1",
        "SÍ": "1",
        "NO": "2",
    },
    "EXPERIENCIA_LABORAL": {
        "SI": "1",
        "SÍ": "1",
        "NO": "2",
    },
    "VIGENCIA_CARRERA": {
        "VIGENTE CON ESTUDIANTES NUEVOS": "1",
        "VIGENTE SIN ESTUDIANTES NUEVOS": "2",
        "VIGENTE SIN ESTUDIANTES NUEVOS, SOLO QUEDAN ESTUDIANTES ANTIGUOS": "2",
        "NO VIGENTE": "3",
        "NO VIGENTE, SIN ESTUDIANTES": "3",
    },
}


def norm(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text)
    return text.upper()


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def convert_value(field: str, raw: Any) -> tuple[str, str | None]:
    value = text_value(raw)
    if value == "":
        return "", None

    if field in MAPAS:
        key = norm(value)
        mapping = {norm(k): v for k, v in MAPAS[field].items()}
        if key in mapping:
            return mapping[key], None
        if re.fullmatch(r"\d+", value):
            return value, None
        return value, f"sin_mapeo_{field}"

    if field == "PROMEDIO_MIN_ENS_MEDIA" and norm(value) in {"NO APLICA", "NO APLICA.", "N/A"}:
        return "0", None

    if field == "FECHA_ADMISION_INICIAL":
        if isinstance(raw, (date, datetime)):
            return raw.strftime("%d/%m/%Y"), None
        return value, None

    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw)), None

    return value, None


def fill_color(cell: Any) -> str:
    fill = cell.fill
    if not fill or not fill.fill_type:
        return ""
    fg = fill.fgColor
    if fg.type == "rgb" and fg.rgb:
        return fg.rgb
    if fg.type == "theme":
        return f"theme:{fg.theme}:tint:{fg.tint}"
    if fg.type == "indexed":
        return f"indexed:{fg.indexed}"
    return str(fg.type)


def read_precarga() -> tuple[list[str], list[dict[str, str]]]:
    with PRECARGA.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=";")
        rows = list(reader)
        return list(reader.fieldnames or []), rows


def read_docencia_modificables() -> tuple[list[str], list[dict[str, Any]], dict[int, dict[str, str]]]:
    wb_values = load_workbook(DOCENCIA_MODIFICABLES, data_only=True)
    ws_values = wb_values[DOCENCIA_HOJA] if DOCENCIA_HOJA else wb_values.active
    headers = [text_value(ws_values.cell(1, c).value) for c in range(1, ws_values.max_column + 1)]

    rows: list[dict[str, Any]] = []
    for r in range(2, ws_values.max_row + 1):
        item = {headers[c - 1]: ws_values.cell(r, c).value for c in range(1, ws_values.max_column + 1)}
        item["_excel_row"] = r
        rows.append(item)

    wb_styles = load_workbook(DOCENCIA_MODIFICABLES, data_only=False)
    ws_styles = wb_styles[DOCENCIA_HOJA] if DOCENCIA_HOJA else wb_styles.active
    styles_by_row: dict[int, dict[str, str]] = defaultdict(dict)
    for r in range(2, ws_styles.max_row + 1):
        for c, header in enumerate(headers, start=1):
            color = fill_color(ws_styles.cell(r, c))
            if color:
                styles_by_row[r][header] = color

    return headers, rows, styles_by_row


def perfil_vacantes() -> dict[str, Any]:
    if not DOCENCIA_VACANTES.exists():
        return {"estado": "no_aplica_en_esta_entrega"}
    wb = load_workbook(DOCENCIA_VACANTES, data_only=True)
    resumen: dict[str, Any] = {}
    for ws in wb.worksheets:
        non_empty_rows = []
        for r in range(1, ws.max_row + 1):
            if any(text_value(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)):
                non_empty_rows.append(r)
        resumen[ws.title] = {
            "filas_no_vacias": len(non_empty_rows),
            "columnas_maximas": ws.max_column,
            "primeras_filas_no_vacias": [
                [text_value(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 15) + 1)]
                for r in non_empty_rows[:8]
            ],
        }
    return resumen


def main() -> None:
    ts = os.environ.get("TS") or datetime.now().strftime("%Y%m%d_%H%M%S")
    for folder in [OUT_TRABAJO, OUT_VALIDACIONES, OUT_PENDIENTES, OUT_LOGS]:
        folder.mkdir(parents=True, exist_ok=True)

    headers_precarga, precarga_rows = read_precarga()
    headers_docencia, docencia_rows, styles_by_row = read_docencia_modificables()

    target_headers = [h for h in headers_precarga if h not in NO_CARGA]

    errores_estructura = []
    encabezados_admitidos = [headers_precarga, [h for h in headers_precarga if h != "VERSION"]]
    if headers_docencia not in encabezados_admitidos:
        errores_estructura.append("Los encabezados de Docencia no coinciden con la precarga SIES, completa o sin VERSION.")
    if len(precarga_rows) != len(docencia_rows):
        errores_estructura.append(f"Cantidad de registros distinta: precarga={len(precarga_rows)} docencia={len(docencia_rows)}.")
    if len(target_headers) != 48:
        errores_estructura.append(f"La estructura de salida no tiene 48 columnas: tiene {len(target_headers)}.")

    cambios: list[dict[str, Any]] = []
    problemas_conversion: list[dict[str, Any]] = []
    filas_naranjas: list[dict[str, Any]] = []
    output_rows: list[dict[str, str]] = []

    orange_color = "theme:5:tint:0.5999938962981048"

    for idx, (base, doc) in enumerate(zip(precarga_rows, docencia_rows), start=2):
        for key in LLAVE_ORDEN:
            if key not in headers_docencia:
                continue
            if text_value(base.get(key)) != text_value(doc.get(key)):
                errores_estructura.append(
                    f"Diferencia de llave/orden fila {idx}, campo {key}: precarga={base.get(key)!r}, docencia={doc.get(key)!r}"
                )

        row = dict(base)
        for field in CAMPOS_MODIFICABLES:
            raw = doc.get(field)
            converted, problem = convert_value(field, raw)
            if converted != text_value(base.get(field)):
                cambios.append({
                    "fila_excel": idx,
                    "cod_carrera": text_value(doc.get("COD_CARRERA")),
                    "nombre_carrera": text_value(doc.get("NOMBRE_CARRERA")),
                    "modalidad_docencia": text_value(doc.get("MODALIDAD")),
                    "jornada_docencia": text_value(doc.get("COD_JORNADA")),
                    "version": text_value(doc.get("VERSION") or base.get("VERSION")),
                    "campo": field,
                    "valor_precarga": text_value(base.get(field)),
                    "valor_docencia": text_value(raw),
                    "valor_salida": converted,
                })
            if problem:
                problemas_conversion.append({
                    "fila_excel": idx,
                    "cod_carrera": text_value(doc.get("COD_CARRERA")),
                    "nombre_carrera": text_value(doc.get("NOMBRE_CARRERA")),
                    "campo": field,
                    "valor": text_value(raw),
                    "problema": problem,
                })
            row[field] = converted

        orange_fields = [field for field, color in styles_by_row.get(idx, {}).items() if color == orange_color]
        if orange_fields:
            filas_naranjas.append({
                "fila_excel": idx,
                "cod_carrera": text_value(doc.get("COD_CARRERA")),
                "nombre_carrera": text_value(doc.get("NOMBRE_CARRERA")),
                "modalidad_docencia": text_value(doc.get("MODALIDAD")),
                "jornada_docencia": text_value(doc.get("COD_JORNADA")),
                    "version": text_value(doc.get("VERSION") or base.get("VERSION")),
                "campos_naranjo": orange_fields,
                "vigencia_salida": row.get("VIGENCIA_CARRERA", ""),
                "vacantes_primer_semestre_salida": row.get("VACANTES_PRIMER_SEMESTRE", ""),
                "vacantes_segundo_semestre_salida": row.get("VACANTES_SEGUNDO_SEMESTRE", ""),
            })

        output_rows.append({h: text_value(row.get(h, "")) for h in target_headers})

    faltantes: list[dict[str, Any]] = []
    for i, row in enumerate(output_rows, start=2):
        if row.get("VIGENCIA_CARRERA") != "1":
            continue
        missing = [field for field in OBLIGATORIOS_VIGENCIA_1 if text_value(row.get(field)) == ""]
        if missing:
            faltantes.append({
                "fila_prefinal": i,
                "cod_carrera": row.get("COD_CARRERA", ""),
                "nombre_carrera": row.get("NOMBRE_CARRERA", ""),
                "modalidad": row.get("MODALIDAD", ""),
                "cod_jornada": row.get("COD_JORNADA", ""),
                "version": row.get("VERSION", ""),
                "faltantes": missing,
            })

    resumen_faltantes = Counter()
    for item in faltantes:
        resumen_faltantes.update(item["faltantes"])

    carga_bloqueada = bool(errores_estructura or problemas_conversion or faltantes or filas_naranjas)

    prefinal = OUT_TRABAJO / f"PREFINAL_ETAPA1_DOCENCIA02_48_COLUMNAS_CON_ENCABEZADO_{ts}.csv"
    with prefinal.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=target_headers, delimiter=";")
        writer.writeheader()
        writer.writerows(output_rows)

    cambios_path = OUT_VALIDACIONES / f"CAMBIOS_DOCENCIA02_VS_PRECARGA_SIES_{ts}.tsv"
    with cambios_path.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "fila_excel",
            "cod_carrera",
            "nombre_carrera",
            "modalidad_docencia",
            "jornada_docencia",
            "version",
            "campo",
            "valor_precarga",
            "valor_docencia",
            "valor_salida",
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(cambios)

    pendientes_path = OUT_PENDIENTES / f"PENDIENTES_BLOQUEANTES_DOCENCIA02_ETAPA1_{ts}.tsv"
    with pendientes_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter="\t", lineterminator="\n")
        writer.writerow(["tipo", "fila", "cod_carrera", "nombre_carrera", "modalidad", "jornada", "version", "detalle"])
        for item in faltantes:
            writer.writerow([
                "FALTANTE_OBLIGATORIO_VIGENCIA_1",
                item["fila_prefinal"],
                item["cod_carrera"],
                item["nombre_carrera"],
                item["modalidad"],
                item["cod_jornada"],
                item["version"],
                ",".join(item["faltantes"]),
            ])
        for item in filas_naranjas:
            writer.writerow([
                "REVISION_NARANJO_RAP_VACANTES",
                item["fila_excel"],
                item["cod_carrera"],
                item["nombre_carrera"],
                item["modalidad_docencia"],
                item["jornada_docencia"],
                item["version"],
                ",".join(item["campos_naranjo"]),
            ])
        for item in problemas_conversion:
            writer.writerow([
                "PROBLEMA_CONVERSION",
                item["fila_excel"],
                item["cod_carrera"],
                item["nombre_carrera"],
                "",
                "",
                "",
                f"{item['campo']}={item['valor']} ({item['problema']})",
            ])
        for error in errores_estructura:
            writer.writerow(["ERROR_ESTRUCTURA", "", "", "", "", "", "", error])

    resumen = {
        "proceso": "SIES Oferta Academica-Acceso 2027",
        "etapa": "Etapa 1 - Oferta Academica Vigente Editada",
        "timestamp": ts,
        "fuentes": {
            "precarga_sies": str(PRECARGA),
            "docencia_modificables": str(DOCENCIA_MODIFICABLES),
            "docencia_vacantes": str(DOCENCIA_VACANTES),
        },
        "salidas": {
            "prefinal_revision_48_columnas_con_encabezado": str(prefinal),
            "cambios_docencia_vs_precarga": str(cambios_path),
            "pendientes_bloqueantes": str(pendientes_path),
        },
        "criterio": {
            "base": "Precarga SIES original.",
            "sobrescritura": "Solo campos modificables de Etapa 1 desde Docencia 02, convertidos a codigos PES cuando corresponde.",
            "exclusiones_carga": sorted(NO_CARGA),
            "advertencia": "No se genera CSV final sin encabezado mientras existan bloqueos.",
        },
        "totales": {
            "registros_precarga": len(precarga_rows),
            "registros_docencia": len(docencia_rows),
            "columnas_precarga": len(headers_precarga),
            "columnas_prefinal": len(target_headers),
            "cambios_detectados": len(cambios),
            "filas_con_faltantes_obligatorios_vigencia_1": len(faltantes),
            "campos_faltantes_vigencia_1": dict(sorted(resumen_faltantes.items())),
            "filas_naranjas_revision": len(filas_naranjas),
            "problemas_conversion": len(problemas_conversion),
            "errores_estructura": len(errores_estructura),
            "carga_bloqueada": carga_bloqueada,
            "vigencia_salida": dict(sorted(Counter(row.get("VIGENCIA_CARRERA", "") for row in output_rows).items())),
        },
        "vacantes_segundo_excel_perfil": perfil_vacantes(),
    }

    resumen_path = OUT_VALIDACIONES / f"RESUMEN_INTEGRACION_DOCENCIA02_ETAPA1_{ts}.json"
    resumen_path.write_text(json.dumps(resumen, ensure_ascii=False, indent=2), encoding="utf-8")

    log_path = OUT_LOGS / f"integracion_docencia02_etapa1_{ts}.log"
    log_path.write_text(
        "\n".join([
            f"prefinal={prefinal}",
            f"resumen={resumen_path}",
            f"pendientes={pendientes_path}",
            f"cambios={cambios_path}",
            f"carga_bloqueada={carga_bloqueada}",
        ]) + "\n",
        encoding="utf-8",
    )

    print(json.dumps({
        "prefinal": str(prefinal),
        "resumen": str(resumen_path),
        "pendientes": str(pendientes_path),
        "cambios": str(cambios_path),
        "totales": resumen["totales"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
