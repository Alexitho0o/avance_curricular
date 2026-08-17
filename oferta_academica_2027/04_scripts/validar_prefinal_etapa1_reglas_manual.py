#!/usr/bin/env python3
from __future__ import annotations

import csv
import os
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PREFINAL = Path(os.environ.get(
    "PREFINAL_ETAPA1",
    ROOT / "05_datos_trabajo" / "PREFINAL_ETAPA1_DOCENCIA02_48_COLUMNAS_CON_ENCABEZADO_20260814_081424_v2.csv",
))
PRECARGA = ROOT / "02_precarga_pes" / "reporte_dinamico_5913_2026-08-12_08:22:23.csv"
DOCENCIA_MODIFICABLES = Path(os.environ.get(
    "DOCENCIA_MODIFICABLES",
    ROOT / "03_fuentes_institucionales" / "entregas_docencia" / "20260814_entrega_02" / "reporte_dinamico_5913_2026-08-12_14_52_50 Modificables_lm.xlsx",
))
DOCENCIA_HOJA = os.environ.get("DOCENCIA_HOJA", "")
REGLAS = ROOT / "11_gobernanza" / "reglas_validacion_etapa1_manual" / "REGLAS_MANUAL_ETAPA1_OFERTA_ACADEMICA_2027.json"
OUT_DIR = ROOT / "06_validaciones"

NO_CARGA = {"CODIGO_IES_NUM", "CANTIDAD_MATRICULA_DFE", "CANTIDAD_BENEFICIO_DFE"}
AREAS = [
    "AREA_ADMIN_DERECHO",
    "AREA_AGRI_SILVI_PESCA_VET",
    "AREA_ARTES_HUMANIDADES",
    "AREA_CIENCIAS_NAT_MAT_ESTAD",
    "AREA_CS_SOCIAL_PERIODISMO_INFO",
    "AREA_EDUCACION",
    "AREA_INGE_INDUSTRIA_CONSTRUC",
    "AREA_SALUD_BIENESTAR",
    "AREA_SERVICIOS",
    "AREA_TECNO_INFO_COMUNICA",
]

AREA_BY_ACTUAL_NIVEL_1 = {
    "1": "AREA_ADMIN_DERECHO",
    "2": "AREA_AGRI_SILVI_PESCA_VET",
    "3": "AREA_ARTES_HUMANIDADES",
    "4": "AREA_CIENCIAS_NAT_MAT_ESTAD",
    "5": "AREA_CS_SOCIAL_PERIODISMO_INFO",
    "6": "AREA_EDUCACION",
    "7": "AREA_INGE_INDUSTRIA_CONSTRUC",
    "8": "AREA_SALUD_BIENESTAR",
    "9": "AREA_SERVICIOS",
    "10": "AREA_TECNO_INFO_COMUNICA",
}

MODIFICABLES = {
    "REGIMEN",
    "DURACION_REGIMEN",
    "ACREDITACION",
    "REQUISITO_INGRESO",
    "SEMESTRES_RECONOCIDOS",
    "AREA_ADMIN_DERECHO",
    "AREA_AGRI_SILVI_PESCA_VET",
    "AREA_ARTES_HUMANIDADES",
    "AREA_CIENCIAS_NAT_MAT_ESTAD",
    "AREA_CS_SOCIAL_PERIODISMO_INFO",
    "AREA_EDUCACION",
    "AREA_INGE_INDUSTRIA_CONSTRUC",
    "AREA_SALUD_BIENESTAR",
    "AREA_SERVICIOS",
    "AREA_TECNO_INFO_COMUNICA",
    "VACANTES_PRIMER_SEMESTRE",
    "VACANTES_SEGUNDO_SEMESTRE",
    "FECHA_ADMISION_INICIAL",
    "ENLACE_INFO_PROGRAMA",
    "LICENCIA_ENS_MEDIA",
    "NOTAS_ENS_MEDIA",
    "PROMEDIO_MIN_ENS_MEDIA",
    "RECONOCIMIENTOS_APREN_PREVIOS",
    "EXPERIENCIA_LABORAL",
    "MAIL_DIFUSION_CARRERA",
    "FORMATO_VALOR",
    "VALOR_MATRICULA_ANUAL",
    "COSTO_TITULACION",
    "VALOR_CERTIFICADO_DIPLOMA",
    "ARANCEL_ANUAL",
    "VIGENCIA_CARRERA",
}


def s(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def n(value: Any) -> int | None:
    text = s(value)
    if text == "":
        return None
    try:
        return int(float(text.replace(",", ".")))
    except ValueError:
        return None


def dec(value: Any) -> float | None:
    text = s(value).replace(",", ".")
    if text == "":
        return None
    try:
        return float(text)
    except ValueError:
        return None


def read_csv(path: Path, delimiter: str) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        return list(reader.fieldnames or []), list(reader)


def issue(issues: list[dict[str, Any]], rule: str, severity: str, row: int | str, record: dict[str, str], field: str, message: str) -> None:
    issues.append({
        "regla_id": rule,
        "severidad": severity,
        "fila": row,
        "cod_carrera": record.get("COD_CARRERA", ""),
        "nombre_carrera": record.get("NOMBRE_CARRERA", ""),
        "modalidad": record.get("MODALIDAD", ""),
        "cod_jornada": record.get("COD_JORNADA", ""),
        "version": record.get("VERSION", ""),
        "campo": field,
        "valor": record.get(field, ""),
        "detalle": message,
    })


def load_orange_rows() -> dict[int, list[str]]:
    if not DOCENCIA_MODIFICABLES.exists():
        return {}
    wb = load_workbook(DOCENCIA_MODIFICABLES, data_only=False)
    ws = wb[DOCENCIA_HOJA] if DOCENCIA_HOJA else wb.active
    headers = [s(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    orange = "theme:5:tint:0.5999938962981048"
    out: dict[int, list[str]] = defaultdict(list)
    for r in range(2, ws.max_row + 1):
        for c, field in enumerate(headers, start=1):
            fill = ws.cell(r, c).fill
            if not fill or not fill.fill_type:
                continue
            fg = fill.fgColor
            color = ""
            if fg.type == "theme":
                color = f"theme:{fg.theme}:tint:{fg.tint}"
            elif fg.type == "rgb":
                color = fg.rgb or ""
            if color == orange:
                out[r].append(field)
    return dict(out)


def validate() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    headers, rows = read_csv(PREFINAL, ";")
    headers_precarga, precarga = read_csv(PRECARGA, ";")
    rules_data = json.loads(REGLAS.read_text(encoding="utf-8"))

    issues: list[dict[str, Any]] = []
    orange_rows = load_orange_rows()

    if len(headers) != 48:
        issue(issues, "OA1-001", "BLOQUEANTE", "archivo", {}, "estructura", f"El prefinal tiene {len(headers)} columnas; debe tener 48.")
    for forbidden in sorted(NO_CARGA):
        if forbidden in headers:
            issue(issues, "OA1-001", "BLOQUEANTE", "archivo", {}, forbidden, "Columna prohibida presente en archivo de carga Etapa 1.")
    if len(rows) != len(precarga):
        issue(issues, "OA1-002", "BLOQUEANTE", "archivo", {}, "registros", f"Prefinal={len(rows)} registros; precarga={len(precarga)} registros.")

    target_precarga_headers = [h for h in headers_precarga if h not in NO_CARGA]
    if headers != target_precarga_headers:
        issue(issues, "OA1-001", "BLOQUEANTE", "archivo", {}, "estructura", "El orden de columnas no coincide con la precarga sin las 3 columnas excluidas.")

    for idx, row in enumerate(rows, start=2):
        base = precarga[idx - 2] if idx - 2 < len(precarga) else {}

        for key in ["COD_SEDE", "COD_CARRERA", "NOMBRE_CARRERA", "VERSION"]:
            if s(row.get(key)) != s(base.get(key)):
                issue(issues, "OA1-002", "BLOQUEANTE", idx, row, key, f"No coincide con precarga: {base.get(key)!r}.")

        for field in headers:
            if field in MODIFICABLES:
                continue
            if s(row.get(field)) != s(base.get(field)):
                issue(issues, "OA1-003", "BLOQUEANTE", idx, row, field, f"Campo no modificable distinto a precarga: {base.get(field)!r}.")

        if row.get("COD_NIVEL_GLOBAL") != "1":
            issue(issues, "OA1-004", "BLOQUEANTE", idx, row, "COD_NIVEL_GLOBAL", "Debe ser 1 Pregrado.")
        if row.get("COD_NIVEL_CARRERA") not in {"0", "1", "2", "3", "4"}:
            issue(issues, "OA1-004", "BLOQUEANTE", idx, row, "COD_NIVEL_CARRERA", "Debe estar entre 0 y 4.")

        modalidad = row.get("MODALIDAD", "")
        jornada = row.get("COD_JORNADA", "")
        if jornada in {"1", "2"} and modalidad == "3":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "MODALIDAD", "Diurna/Vespertina no puede ser No Presencial.")
        if jornada == "3" and modalidad != "2":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "MODALIDAD", "Jornada semipresencial exige modalidad 2.")
        if jornada == "4" and modalidad != "3":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "MODALIDAD", "Jornada a distancia exige modalidad 3.")
        if jornada == "5" and modalidad == "3":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "MODALIDAD", "Jornada otro no puede ser no presencial.")
        if modalidad == "1" and jornada not in {"1", "2", "5"}:
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "COD_JORNADA", "Modalidad presencial solo permite jornadas 1, 2 o 5.")
        if modalidad == "2" and jornada == "4":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "COD_JORNADA", "Modalidad semipresencial no puede ser a distancia.")
        if modalidad == "3" and jornada != "4":
            issue(issues, "OA1-005", "BLOQUEANTE", idx, row, "COD_JORNADA", "Modalidad no presencial exige jornada 4.")

        domains = {
            "MODALIDAD": {"1", "2", "3"},
            "COD_JORNADA": {"1", "2", "3", "4", "5"},
            "COD_TIPO_PLAN_CARRERA": {"1", "2", "3"},
            "REGIMEN": {"1", "2", "3", "4"},
            "ACREDITACION": {"1", "2"},
            "FORMATO_VALOR": {"1", "2", ""},
            "VIGENCIA_CARRERA": {"1", "2", "3"},
        }
        for field, allowed in domains.items():
            if row.get(field, "") not in allowed:
                issue(issues, "OA1-006", "BLOQUEANTE", idx, row, field, f"Valor fuera de catálogo permitido: {sorted(allowed)}.")
        req = n(row.get("REQUISITO_INGRESO"))
        if req is None or not (1 <= req <= 10):
            issue(issues, "OA1-006", "BLOQUEANTE", idx, row, "REQUISITO_INGRESO", "Debe estar entre 1 y 10.")

        area_actual = n(row.get("AREA_ACTUAL"))
        if area_actual is None or not (1 <= area_actual <= 10):
            issue(issues, "OA1-007", "BLOQUEANTE", idx, row, "AREA_ACTUAL", "Debe estar entre 1 y 10.")
        for area in AREAS:
            if row.get(area) not in {"0", "1"}:
                issue(issues, "OA1-007", "BLOQUEANTE", idx, row, area, "Área destino debe ser 0 o 1.")

        nivel = row.get("COD_NIVEL_CARRERA")
        area_sum = sum(1 for area in AREAS if row.get(area) == "1")
        if nivel == "1":
            expected = AREA_BY_ACTUAL_NIVEL_1.get(row.get("AREA_ACTUAL", ""))
            if expected and row.get(expected) != "1":
                issue(issues, "OA1-008", "BLOQUEANTE", idx, row, expected, "TNS debe marcar al menos el área destino correspondiente a AREA_ACTUAL.")
            if area_sum > 5:
                issue(issues, "OA1-008", "BLOQUEANTE", idx, row, "areas_destino", "TNS no puede asignar más de 5 áreas destino.")
        if nivel in {"2", "4"} and area_sum != 0:
            issue(issues, "OA1-008", "BLOQUEANTE", idx, row, "areas_destino", "Nivel 2 o 4 debe tener áreas destino en 0.")
        if nivel == "3":
            for area in AREAS:
                expected = "1" if area == "AREA_EDUCACION" else "0"
                if row.get(area) != expected:
                    issue(issues, "OA1-008", "BLOQUEANTE", idx, row, area, "Nivel 3 exige AREA_EDUCACION=1 y resto 0.")

        if row.get("COD_TIPO_PLAN_CARRERA") == "1" and row.get("SEMESTRES_RECONOCIDOS") != "0":
            issue(issues, "OA1-009", "BLOQUEANTE", idx, row, "SEMESTRES_RECONOCIDOS", "Plan regular exige semestres reconocidos igual a 0.")

        de = n(row.get("DURACION_ESTUDIOS"))
        dt = n(row.get("DURACION_TITULACION"))
        df = n(row.get("DURACION_TOTAL"))
        dr = n(row.get("DURACION_REGIMEN"))
        if de is None or de == 0 or de > 24:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_ESTUDIOS", "Debe ser >0 y <=24.")
        if de is not None and de > 14:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_ESTUDIOS", "Manual indica DURACION_ESTUDIOS <=14 o contactar encargado SIES.")
        if dt is None or dt < 0 or dt > 24:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_TITULACION", "Debe ser >=0 y <=24.")
        if df is None or df == 0 or df > 24:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_TOTAL", "Debe ser >0 y <=24.")
        if dr is None or dr <= 0:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_REGIMEN", "Debe ser mayor a 0 y no nulo.")
        if de is not None and df is not None and df < de:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_TOTAL", "No puede ser menor a DURACION_ESTUDIOS.")
        if de is not None and dt is not None and df is not None and de + dt < df:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_TOTAL", "DURACION_ESTUDIOS + DURACION_TITULACION no puede ser menor a DURACION_TOTAL.")
        if row.get("REGIMEN") == "1" and dr is not None and df is not None and dr != df:
            issue(issues, "OA1-010", "BLOQUEANTE", idx, row, "DURACION_REGIMEN", "Si REGIMEN=1, DURACION_REGIMEN debe ser igual a DURACION_TOTAL.")

        vig = row.get("VIGENCIA_CARRERA")
        if vig == "1":
            required = [
                "ARANCEL_ANUAL",
                "COSTO_TITULACION",
                "FORMATO_VALOR",
                "RECONOCIMIENTOS_APREN_PREVIOS",
                "VALOR_CERTIFICADO_DIPLOMA",
                "VALOR_MATRICULA_ANUAL",
                "EXPERIENCIA_LABORAL",
                "FECHA_ADMISION_INICIAL",
                "ENLACE_INFO_PROGRAMA",
                "LICENCIA_ENS_MEDIA",
                "MAIL_DIFUSION_CARRERA",
                "NOTAS_ENS_MEDIA",
                "PROMEDIO_MIN_ENS_MEDIA",
            ]
            for field in required:
                if s(row.get(field)) == "":
                    issue(issues, "OA1-011", "BLOQUEANTE", idx, row, field, "Campo obligatorio cuando VIGENCIA_CARRERA=1.")
            vac1 = n(row.get("VACANTES_PRIMER_SEMESTRE"))
            vac2 = n(row.get("VACANTES_SEGUNDO_SEMESTRE"))
            if vac1 is None:
                issue(issues, "OA1-012", "BLOQUEANTE", idx, row, "VACANTES_PRIMER_SEMESTRE", "No puede ser nulo cuando VIGENCIA_CARRERA=1.")
            elif vac1 == 0:
                issue(issues, "OA1-012", "BLOQUEANTE", idx, row, "VACANTES_PRIMER_SEMESTRE", "Manual indica que vacantes semestrales no pueden ser 0 con vigencia 1.")
            if vac2 is None:
                issue(issues, "OA1-012", "BLOQUEANTE", idx, row, "VACANTES_SEGUNDO_SEMESTRE", "No puede ser nulo cuando VIGENCIA_CARRERA=1.")
            elif vac2 == 0:
                issue(issues, "OA1-012", "BLOQUEANTE", idx, row, "VACANTES_SEGUNDO_SEMESTRE", "Manual indica que vacantes semestrales no pueden ser 0 con vigencia 1.")
        if vig == "2":
            if s(row.get("FECHA_ADMISION_INICIAL")) != "":
                issue(issues, "OA1-013", "BLOQUEANTE", idx, row, "FECHA_ADMISION_INICIAL", "Debe quedar vacío con VIGENCIA_CARRERA=2.")
            for field in ["VACANTES_PRIMER_SEMESTRE", "VACANTES_SEGUNDO_SEMESTRE"]:
                if n(row.get(field)) not in {0, None} or s(row.get(field)) == "":
                    issue(issues, "OA1-013", "BLOQUEANTE", idx, row, field, "Debe ser 0 con VIGENCIA_CARRERA=2.")

        notas = row.get("NOTAS_ENS_MEDIA")
        promedio = dec(row.get("PROMEDIO_MIN_ENS_MEDIA"))
        if notas not in {"1", "2", ""}:
            issue(issues, "OA1-014", "BLOQUEANTE", idx, row, "NOTAS_ENS_MEDIA", "Debe ser 1/SI o 2/NO.")
        if notas == "2" and promedio != 0:
            issue(issues, "OA1-014", "BLOQUEANTE", idx, row, "PROMEDIO_MIN_ENS_MEDIA", "Si NOTAS_ENS_MEDIA=NO, promedio debe ser 0.")
        if notas == "1" and (promedio is None or promedio < 4 or promedio > 7):
            issue(issues, "OA1-014", "BLOQUEANTE", idx, row, "PROMEDIO_MIN_ENS_MEDIA", "Si NOTAS_ENS_MEDIA=SI, promedio debe estar entre 4,00 y 7,00.")

        if vig == "1":
            url = row.get("ENLACE_INFO_PROGRAMA", "")
            if url and not re.match(r"^(https?://|www\.)\S+", url, flags=re.IGNORECASE):
                issue(issues, "OA1-015", "BLOQUEANTE", idx, row, "ENLACE_INFO_PROGRAMA", "Debe ser link completo http(s) o www.")
        mail = row.get("MAIL_DIFUSION_CARRERA", "")
        if mail and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", mail):
            issue(issues, "OA1-015", "BLOQUEANTE", idx, row, "MAIL_DIFUSION_CARRERA", "Formato de correo inválido.")

        for field in ["ARANCEL_ANUAL", "VALOR_MATRICULA_ANUAL"]:
            value = dec(row.get(field))
            if value is not None and value < 0 and value != -1:
                issue(issues, "OA1-016", "BLOQUEANTE", idx, row, field, "Debe ser >=0 o -1.")
        for field in ["COSTO_TITULACION", "VALOR_CERTIFICADO_DIPLOMA"]:
            value = dec(row.get(field))
            if value is not None and value < 0:
                issue(issues, "OA1-016", "BLOQUEANTE", idx, row, field, "Debe ser >=0.")

        # Date validation only runs after the currently missing field is supplied.
        fecha = row.get("FECHA_ADMISION_INICIAL", "")
        if fecha:
            match = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", fecha) or re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", fecha)
            if not match:
                issue(issues, "OA1-017", "BLOQUEANTE", idx, row, "FECHA_ADMISION_INICIAL", "Formato de fecha no reconocido; usar dd/mm/aaaa o aaaa-mm-dd.")
            else:
                if "-" in fecha:
                    yyyy, mm, dd = map(int, match.groups())
                else:
                    dd, mm, yyyy = map(int, match.groups())
                num = yyyy * 10000 + mm * 100 + dd
                if num < 20251002 or num > 20260404:
                    issue(issues, "OA1-017", "BLOQUEANTE", idx, row, "FECHA_ADMISION_INICIAL", "Fuera de rango manual 02/10/2025 a 04/04/2026.")

        anio = n(row.get("ANIO_INICIO"))
        if anio is None or anio > 2025:
            issue(issues, "OA1-018", "BLOQUEANTE", idx, row, "ANIO_INICIO", "No puede ser posterior a 2025 para Vigente Editada.")

        if row.get("VIGENCIA_CARRERA") == "3":
            mat = n(base.get("CANTIDAD_MATRICULA_DFE")) or 0
            ben = n(base.get("CANTIDAD_BENEFICIO_DFE")) or 0
            if mat > 0 or ben > 0:
                issue(issues, "OA1-019", "BLOQUEANTE", idx, row, "VIGENCIA_CARRERA", f"Programa protegido DFE: matricula={mat}, beneficio={ben}.")

        if idx in orange_rows:
            issue(issues, "OA1-021", "REVISION", idx, row, ",".join(orange_rows[idx]), "Fila marcada en naranjo por Docencia: requiere confirmación RAP/vacantes online.")

    # Manual mentions fields not present in official 48-column structure; keep visible but not as row failure.
    for field in ["MALLA_CURRICULAR", "PERFIL_EGRESO"]:
        if field not in headers:
            issue(issues, "OA1-020", "REVISION", "archivo", {}, field, "Campo mencionado en bloque de errores, pero ausente en estructura oficial Etapa 1 de 48 columnas.")

    summary = {
        "reglas_fuente": str(REGLAS),
        "prefinal": str(PREFINAL),
        "precarga": str(PRECARGA),
        "total_registros": len(rows),
        "total_columnas": len(headers),
        "issues_total": len(issues),
        "por_severidad": dict(Counter(i["severidad"] for i in issues)),
        "por_regla": dict(sorted(Counter(i["regla_id"] for i in issues).items())),
        "por_campo": dict(sorted(Counter(i["campo"] for i in issues).items())),
        "carga_bloqueada": any(i["severidad"] == "BLOQUEANTE" for i in issues),
        "reglas_catalogadas": len(json.loads(REGLAS.read_text(encoding="utf-8"))["reglas"]),
    }
    return issues, summary


def main() -> None:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    issues, summary = validate()

    json_path = OUT_DIR / f"VALIDACION_REGLAS_MANUAL_ETAPA1_DOCENCIA02_{ts}.json"
    tsv_path = OUT_DIR / f"VALIDACION_REGLAS_MANUAL_ETAPA1_DOCENCIA02_{ts}.tsv"
    md_path = OUT_DIR / f"VALIDACION_REGLAS_MANUAL_ETAPA1_DOCENCIA02_{ts}.md"

    json_path.write_text(json.dumps({"timestamp": ts, "resumen": summary, "hallazgos": issues}, ensure_ascii=False, indent=2), encoding="utf-8")

    with tsv_path.open("w", encoding="utf-8", newline="") as f:
        fieldnames = ["regla_id", "severidad", "fila", "cod_carrera", "nombre_carrera", "modalidad", "cod_jornada", "version", "campo", "valor", "detalle"]
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(issues)

    lines = [
        "# Validación Reglas Manual Etapa 1 - Docencia 02",
        "",
        f"Prefinal: {PREFINAL}",
        f"Reglas: {REGLAS}",
        "",
        "## Resumen",
        "",
    ]
    for key, value in summary.items():
        lines.append(f"- {key}: {value}")
    lines.extend(["", "## Primeros Hallazgos", "", "| Regla | Severidad | Fila | Carrera | Campo | Detalle |", "|---|---|---|---|---|---|"])
    for item in issues[:80]:
        lines.append(f"| {item['regla_id']} | {item['severidad']} | {item['fila']} | {item['nombre_carrera']} | {item['campo']} | {item['detalle']} |")
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(json.dumps({
        "json": str(json_path),
        "tsv": str(tsv_path),
        "md": str(md_path),
        "resumen": summary,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
