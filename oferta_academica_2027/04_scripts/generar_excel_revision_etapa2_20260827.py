# -*- coding: utf-8 -*-
# Fase 6: Excel de revision para Docencia - Etapa 2
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.expanduser("~/mnt/avance_curricular/oferta_academica_2027")
VAL_DIR = os.path.join(BASE, "06_validaciones")
INST = os.path.join(BASE, "03_fuentes_institucionales/entregas_docencia/20260827_etapa2_entrega_01/Oferta 2027 Etapa2_prueba_1.xlsx")
STAMP = "20260827"
OUT_DIR = os.path.join(BASE, "07_resultados", "reportes_pendientes_etapa2")
os.makedirs(OUT_DIR, exist_ok=True)

val = json.load(open(os.path.join(VAL_DIR, f"VALIDACION_COMPLETA_ETAPA2_{STAMP}.json"), encoding="utf-8"))
concil = json.load(open(os.path.join(VAL_DIR, f"CONCILIACION_ETAPA2_VS_REPORTE_5912_{STAMP}.json"), encoding="utf-8"))
contrato = json.load(open(os.path.join(BASE, "11_gobernanza", f"REGLAS_MANUAL_ETAPA2_OFERTA_NUEVA_2027_{STAMP}.json"), encoding="utf-8"))

RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="C00000")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = openpyxl.Workbook()

# ---------- Hoja Resumen ----------
ws1 = wb.active
ws1.title = "Resumen"
ws1["A1"] = "OFERTA ACADEMICA NUEVA 2027 - ETAPA 2 (TP Adscritas) - REPORTE DE PENDIENTES"
ws1["A1"].font = Font(bold=True, size=14)
ws1["A2"] = "Generado: 2026-08-27  |  Archivo institucional: Oferta 2027 Etapa2_prueba_1.xlsx (copia gobernada 20260827_etapa2_entrega_01)"
ws1["A4"] = "ESTADO: NO CARGAR A SIES"
ws1["A4"].font = TITLE_FONT
ws1["A5"] = "Existen hallazgos BLOQUEANTES en el 100% de los registros (50/50) y contradicciones documentales no resueltas en el propio instructivo (VERSION, rango de fechas 2027, MALLA_CURRICULAR/PERFIL_EGRESO). No se genera CSV de carga."
ws1["A5"].alignment = Alignment(wrap_text=True)
ws1.merge_cells("A5:F5")
ws1.row_dimensions[5].height = 45

r = 7
ws1.cell(row=r, column=1, value="Metrica").font = BOLD
ws1.cell(row=r, column=2, value="Valor").font = BOLD
r += 1
metricas = [
    ("Filas totales en el libro (incluye encabezado)", 55),
    ("Registros de oferta detectados (filas reales de datos)", val["total_filas_evaluadas"]),
    ("Columnas del archivo recibido", 49),
    ("Columnas oficiales de Etapa 2 (Anexo2, pag.35-41)", 48),
    ("Estructura oficial", "PENDIENTE DE CONFIRMACION (ver hoja Reglas aplicadas)"),
    ("Total hallazgos de validacion", len(val["hallazgos"])),
    ("Hallazgos BLOQUEANTE", val["resumen"].get("BLOQUEANTE", 0)),
    ("Hallazgos REVISION_MANUAL", val["resumen"].get("REVISION_MANUAL", 0)),
    ("Hallazgos INFORMATIVO", val["resumen"].get("INFORMATIVO", 0)),
    ("Filas con al menos un hallazgo BLOQUEANTE", len(set(f["fila_excel"] for f in val["hallazgos"] if f["severidad"]=="BLOQUEANTE"))),
    ("Registros POSIBLE_DUPLICADO_ETAPA1 (coinciden exacto con 5912)", concil["resumen"].get("POSIBLE_DUPLICADO_ETAPA1", 0)),
    ("Registros NUEVO_NO_PRESENTE_EN_ETAPA1", concil["resumen"].get("NUEVO_NO_PRESENTE_EN_ETAPA1", 0)),
    ("Registros REQUIERE_REVISION_MANUAL (conciliacion)", concil["resumen"].get("REQUIERE_REVISION_MANUAL", 0)),
    ("Registros NUEVA_SEDE", concil["resumen"].get("NUEVA_SEDE", 0)),
    ("Registros NUEVA_JORNADA", concil["resumen"].get("NUEVA_JORNADA", 0)),
    ("Registros NUEVO_TIPO_PLAN", concil["resumen"].get("NUEVO_TIPO_PLAN", 0)),
]
for label, value in metricas:
    ws1.cell(row=r, column=1, value=label)
    ws1.cell(row=r, column=2, value=value)
    r += 1

r += 1
ws1.cell(row=r, column=1, value="Contradicciones documentales no resueltas (nivel Instructivo oficial):").font = BOLD
r += 1
for c in contrato["contradicciones_detectadas"]:
    ws1.cell(row=r, column=1, value=f"- {c['campo']}: {c['descripcion']}").alignment = Alignment(wrap_text=True)
    ws1.merge_cells(f"A{r}:F{r}")
    ws1.row_dimensions[r].height = 45
    r += 1

ws1.column_dimensions["A"].width = 55
ws1.column_dimensions["B"].width = 45

# ---------- Hoja Detalle de hallazgos ----------
ws2 = wb.create_sheet("Detalle de hallazgos")
cols2 = ["fila_excel","llave_diagnostica","carrera","sede","modalidad","jornada","campo","valor_observado","regla_incumplida","fuente","pagina","linea","severidad","accion_requerida"]
for j, c in enumerate(cols2, start=1):
    cell = ws2.cell(row=1, column=j, value=c)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
ws2.freeze_panes = "A2"
for i, f in enumerate(val["hallazgos"], start=2):
    for j, c in enumerate(cols2, start=1):
        cell = ws2.cell(row=i, column=j, value=f.get(c))
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in ("regla_incumplida","accion_requerida","llave_diagnostica")))
    if f["severidad"] == "BLOQUEANTE":
        for j in range(1, len(cols2)+1):
            ws2.cell(row=i, column=j).fill = RED
    elif f["severidad"] == "REVISION_MANUAL":
        for j in range(1, len(cols2)+1):
            ws2.cell(row=i, column=j).fill = YELLOW
widths2 = [10,45,30,22,10,8,26,20,55,32,8,8,14,45]
for j, w in enumerate(widths2, start=1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ---------- Hoja Registros (todos los registros del archivo, con conteo de hallazgos) ----------
ws3 = wb.create_sheet("Registros (todos)")
wb_src = openpyxl.load_workbook(INST, data_only=True)
ws_src = wb_src["Hoja4"]
headers_src = [ws_src.cell(row=1, column=c).value for c in range(1, ws_src.max_column+1)]
extra_cols = ["N_BLOQUEANTES", "N_REVISION_MANUAL", "N_INFORMATIVO", "CLASIFICACION_CONCILIACION_5912"]
for j, h in enumerate(headers_src + extra_cols, start=1):
    cell = ws3.cell(row=1, column=j, value=h)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
ws3.freeze_panes = "A2"

count_by_row = {}
for f in val["hallazgos"]:
    d0 = count_by_row.setdefault(f["fila_excel"], {"BLOQUEANTE":0,"REVISION_MANUAL":0,"INFORMATIVO":0})
    d0[f["severidad"]] += 1
concil_by_row = {d["fila_excel"]: d["clasificacion"] for d in concil["detalle"]}

out_r = 2
for r_src in range(2, 56):
    vals = [ws_src.cell(row=r_src, column=c).value for c in range(1, ws_src.max_column+1)]
    if all(v is None for v in vals):
        continue
    for j, v in enumerate(vals, start=1):
        cell = ws3.cell(row=out_r, column=j, value=v)
        cell.border = BORDER
    cnt = count_by_row.get(r_src, {"BLOQUEANTE":0,"REVISION_MANUAL":0,"INFORMATIVO":0})
    ws3.cell(row=out_r, column=len(headers_src)+1, value=cnt["BLOQUEANTE"])
    ws3.cell(row=out_r, column=len(headers_src)+2, value=cnt["REVISION_MANUAL"])
    ws3.cell(row=out_r, column=len(headers_src)+3, value=cnt["INFORMATIVO"])
    ws3.cell(row=out_r, column=len(headers_src)+4, value=concil_by_row.get(r_src, ""))
    if cnt["BLOQUEANTE"] > 0:
        for j in range(1, len(headers_src)+5):
            ws3.cell(row=out_r, column=j).fill = RED
    elif cnt["REVISION_MANUAL"] > 0:
        for j in range(1, len(headers_src)+5):
            ws3.cell(row=out_r, column=j).fill = YELLOW
    out_r += 1
for j in range(1, len(headers_src)+5):
    ws3.column_dimensions[get_column_letter(j)].width = 16

# ---------- Hoja Reglas aplicadas ----------
ws4 = wb.create_sheet("Reglas aplicadas")
cols4 = ["posicion","letra","nombre_oficial","tipo","dominio","obligatorio_o_condicional","permite_vacio","permite_cero","permite_menos1","regla_cruzada","pagina","linea","estado_respaldo"]
for j, c in enumerate(cols4, start=1):
    cell = ws4.cell(row=1, column=j, value=c)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
ws4.freeze_panes = "A2"
for i, col in enumerate(contrato["columnas"], start=2):
    for j, c in enumerate(cols4, start=1):
        cell = ws4.cell(row=i, column=j, value=col.get(c))
        cell.alignment = Alignment(vertical="top", wrap_text=(c in ("dominio","regla_cruzada")))
        cell.border = BORDER
    if col["estado_respaldo"] not in ("RESPALDADO_ESTRUCTURA_MANUAL_PAG35_41",):
        for j in range(1, len(cols4)+1):
            ws4.cell(row=i, column=j).fill = YELLOW
widths4 = [8,6,32,10,45,20,14,14,14,55,8,8,45]
for j, w in enumerate(widths4, start=1):
    ws4.column_dimensions[get_column_letter(j)].width = w

out_path = os.path.join(OUT_DIR, f"REPORTE_PENDIENTES_ETAPA2_OFERTA_NUEVA_2027_{STAMP}.xlsx")
wb.save(out_path)
print("Excel generado:", out_path)
