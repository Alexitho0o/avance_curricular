"""Verify the delivered workbook without modifying it or any source workbook."""

import json
import re
from pathlib import Path
from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import ROOT, FIELDS, sha, write_once


def main():
    tag='maqueta_etapa2_noviembre_20261023_v02'
    audit=ROOT/'06_validaciones'/tag
    data=json.loads((audit/'INSUMOS.json').read_text())
    manifest=json.loads((ROOT/'11_gobernanza'/tag/'MANIFIESTO.json').read_text())
    output=Path(manifest['salidas'][0]['ruta'])
    original=load_workbook(data['source'],data_only=False)
    wb=load_workbook(output,data_only=False)
    cached=load_workbook(output,data_only=True)
    assert wb.sheetnames==['Programas','Pendientes por campo','Diplomados para completar']
    before=original['Programas']
    after=wb['Programas']
    checked=0
    style_changes=[]
    for row in before:
        for cell in row:
            dest=after[cell.coordinate]
            if cell.coordinate not in manifest['ediciones_programas']:
                assert cell.value==dest.value,(cell.coordinate,cell.value,dest.value)
            if cell.has_style:
                a=(cell.font.name,cell.font.sz,cell.font.bold,cell.font.italic,cell.alignment.horizontal,cell.alignment.vertical,cell.alignment.wrap_text,cell.number_format,cell.fill.fgColor.rgb)
                b=(dest.font.name,dest.font.sz,dest.font.bold,dest.font.italic,dest.alignment.horizontal,dest.alignment.vertical,dest.alignment.wrap_text,dest.number_format,dest.fill.fgColor.rgb)
                if a!=b: style_changes.append(cell.coordinate)
            checked+=1
    assert not style_changes,style_changes
    assert after.freeze_panes==before.freeze_panes
    assert after.tables['RevisionProgramas'].ref==before.tables['RevisionProgramas'].ref
    grid=wb['Diplomados para completar']
    values=cached[grid.title]
    assert [grid.cell(8,c).value for c in range(1,49)]==FIELDS
    assert grid.tables['PlantillaDiplomadosEtapa2'].ref=='A8:AV30'
    assert {str(v.sqref) for v in grid.data_validations.dataValidation}=={f'{c}9:{c}30' for c in ['E','F','G','H','M','P','Q','S','T','V','AQ','AV']}
    expected_empty={0,1,4,5,6,7,8,9,10,11,12,13,14,17,18,19,20,21,47}
    assert len(expected_empty)==19
    missing=0
    for i,p in enumerate(data['candidatos'],9):
        for c in range(48):
            v=values.cell(i,c+1).value
            if c in data['blank_indices']:
                assert v is None
            elif c in expected_empty:
                assert v is None
                missing+=1
            elif c==3: assert v==p['nombre_sies_preparado']
            elif c==15: assert v==3
            elif c==16: assert v==5
            elif 22<=c<=31: assert v==0
            elif c in [32,33]: assert v==15
            elif c==42: assert v==1
            elif c in [43,44,45]: assert v==0
            elif c==46: assert v==p['importe']
            else: raise AssertionError(c)
    assert missing==418
    links=0
    for sh in wb:
        for row in sh:
            for c in row:
                if c.hyperlink:
                    links+=1
                    target=c.hyperlink.location
                    match=re.fullmatch(r"'([^']+)'!([A-Z]+[0-9]+)",target)
                    assert match and match[1] in wb.sheetnames,target
                    assert wb[match[1]][match[2]].value is not None,target
                assert not (c.data_type=='f' and 'HYPERLINK' in c.value)
    assert links==192
    errors=[(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type=='e']
    assert not errors,errors
    assert sum(cached['Pendientes por campo'].cell(r,4).value for r in range(9,57))==418
    assert cached['Pendientes por campo']['A4'].value=='Celdas obligatorias vacias: 418. Campos que deben quedar vacios: 9.'
    assert grid['S9'].value is None
    for p,info in data['inputs'].items():
        assert sha(Path(p))==info['sha256']
        assert oct(Path(p).stat().st_mode&0o777)==info['modo']
    assert sha(Path(data['snapshot']))==sha(Path(data['source']))
    for p in manifest['salidas']: assert sha(Path(p['ruta']))==p['sha256']
    report={'estado':'VERIFICADO_PARA_COMPLETAR_NO_PARA_CARGAR','hojas':wb.sheetnames,'celdas_hoja_usuario_verificadas':checked,'cambios_datos_usuario':0,'cambios_estilo_hoja_programas':style_changes,'columnas_sies':48,'programas':22,'campos_obligatorios':39,'campos_vacios':9,'faltantes_reales':missing,'enlaces_nativos_destino_valido':links,'errores_formula':0,'prueba_dinamica_contador':'S9=2 reduce faltantes a 21; restaurar vacio devuelve 22, ejecutado en constructor','originales_y_etapa1_congelada_intactos':True,'notas':sum(bool(c.comment) for s in wb for row in s for c in row),'limitacion':'Enlaces comprobados en OOXML y mediante lectura del archivo; no se opero la interfaz de Microsoft Excel.'}
    write_once(audit/'VERIFICACION_FINAL.json',(json.dumps(report,ensure_ascii=False,indent=2)+'\n').encode())
    print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':
    main()
