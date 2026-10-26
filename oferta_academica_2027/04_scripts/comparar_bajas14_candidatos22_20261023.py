"""Compare names without inferring academic equivalence or changing delivery files."""
import csv
import io
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import ROOT, FIELDS, sha, tsv, write_once


def normalized(name):
    text=''.join(c for c in unicodedata.normalize('NFD',name.upper()) if not unicodedata.combining(c))
    return ' '.join(re.sub(r'[^A-Z0-9 ]',' ',text).split())


def tokens(name):
    stop={'DIPLOMADO','EN','DE','DEL','LA','EL','LOS','LAS','Y','PARA','CON'}
    return tuple(sorted(t for t in normalized(name).split() if t not in stop))


def main():
    frozen=ROOT/'07_resultados/cargas_congeladas/20261023_etapa1_octubre_v06_aceptada'
    csvpath=frozen/'CARGA_OFERTA_2027_ETAPA1_17467_20261023_V06.csv'
    source=Path.home()/'Desktop/DIPLOMADOS_ETAPA2_PARA_COMPLETAR_20261023_V03.xlsx'
    inputs={str(p):{'sha256':sha(p),'modo':oct(p.stat().st_mode&0o777)} for p in [source]+sorted(p for p in frozen.iterdir() if p.is_file())}
    assert sha(csvpath)=='d7fe88035cd1215d448272a9854bee0932c45c200f4f80bb58636b62a6ae4e0b'
    rows=list(csv.reader(io.StringIO(csvpath.read_text()),delimiter=';'))
    assert len(rows)==161 and all(len(r)==48 for r in rows)
    rows=[dict(zip(FIELDS,r)) for r in rows]
    low=[r for r in rows if r['COD_NIVEL_GLOBAL']=='3' and r['COD_NIVEL_CARRERA']=='5' and r['VIGENCIA_CARRERA']=='3']
    assert len(low)==14
    wb=load_workbook(source,data_only=True)
    ws=wb['Diplomados para completar']
    assert [ws.cell(8,c).value for c in range(1,49)]==FIELDS
    candidates={r:ws.cell(r,4).value for r in range(9,31)}
    assert len(candidates)==22 and all(isinstance(v,str) and v for v in candidates.values())
    exact=[];bag=[]
    for a in low:
        for r,b in candidates.items():
            if normalized(a['NOMBRE_CARRERA'])==normalized(b):exact.append((a['COD_CARRERA'],r))
            if tokens(a['NOMBRE_CARRERA'])==tokens(b):bag.append((a['COD_CARRERA'],r))
    assert exact==bag==[], 'Revisar cualquier nueva coincidencia antes de clasificar'
    # Human-reviewed topical leads, explicitly not a name match or institutional equivalence.
    thematic={
        '40':([9,18],'Apoyo docente y docencia clinica tienen objetivos y ambitos distintos a docencia en educacion superior tecnico profesional. Confirmar si alguno es sucesor o si se descontinua.'),
        '95':([r for r,n in candidates.items() if 'INTELIGENCIA ARTIFICIAL' in n],'Ocho propuestas contienen Inteligencia Artificial. La denominacion general anterior no identifica una unica propuesta; no asignar el mismo codigo a varias.'),
        '97':([13],'Comparte logistica/operaciones, pero no acredita supply chain y mineria de requerimientos ni continuidad curricular.'),
        '104':([9],'Herramientas de IA para apoyo docente puede relacionarse con innovacion docente, pero no demuestra que sea el mismo programa.'),
    }
    outside={
        '63':'Infraestructura Cloud (69) ya quedo vigente en Etapa 1 y no forma parte de los 22. Arquitectura Cloud (63) no tiene equivalencia confirmada con 69.',
        '64':'Desarrollo Seguro de Software es de una especialidad distinta; compartir el ambito de software no demuestra Backend.',
        '65':'UX/UI o Desarrollo Seguro de Software no acreditan cobertura de Frontend; no se propone equivalencia por cercania de area.',
        '67':'No hay propuesta explicitamente de Quality Assurance; Desarrollo Seguro de Software no demuestra aseguramiento de calidad.',
        '71':'Ciberseguridad Aplicada (53) y Seguridad Ofensiva (68) ya estan vigentes; no son equivalencias confirmadas de Seguridad en Arquitectura de Redes.',
        '98':'Gestion de Proyectos TI, Scrum o IA para Gestion de Personas no acreditan Habilidades Directivas para Profesionales STEM.',
        '101':'Las propuestas de cuidados o atencion sanitaria no acreditan Salud Familiar con Enfoque Comunitario.',
        '102':'No se observa una propuesta especifica sobre salud y migracion.',
        '103':'No se observa una propuesta especifica sobre violencia de genero en atencion primaria.',
        '105':'Las propuestas de cuidados paliativos o integrales no acreditan Humanizacion en Salud.',
    }
    detail=[]
    for a in sorted(low,key=lambda x:int(x['COD_CARRERA'])):
        code=a['COD_CARRERA'];near,reason=thematic.get(code,([],outside.get(code,'')))
        assert reason
        detail.append([code,a['NOMBRE_CARRERA'],3,'NO CONFIRMADA','AFINIDAD TEMATICA POR REVISAR' if near else 'SIN CANDIDATO ESPECIFICO', ' / '.join(str(r) for r in near),' | '.join(candidates[r] for r in near),reason,'Docencia: confirmar continuidad mediante programa/malla o confirmar que no se ofrece en 2027. No agregar ni reactivar automaticamente.'])
    tag='conciliacion_bajas14_candidatos22_20261023'
    audit=ROOT/'06_validaciones'/tag
    output=ROOT/'10_entrega'/tag/'CRUCE_14_BAJAS_VS_22_CANDIDATOS.tsv'
    tsv(output,[['cod_carrera_etapa1','nombre_etapa1','vigencia_etapa1','representacion_confirmada','resultado','filas_excel_v03','programas_relacionados_no_equivalentes','fundamento','accion']]+detail)
    control={'bajas_etapa1':14,'candidatos_etapa2':22,'coincidencias_nombre_normalizado':0,'coincidencias_tokens_sin_palabras_funcionales':0,'representados_confirmados':0,'con_afinidad_tematica_no_confirmada':4,'sin_candidato_especifico':10,'candidatos_IA_para_codigo95':len(thematic['95'][0]),'cambios_excel':0,'cambios_etapa1':0}
    assert control['candidatos_IA_para_codigo95']==8
    write_once(audit/'CONTROLES.json',(json.dumps(control,ensure_ascii=False,indent=2)+'\n').encode())
    for p,v in inputs.items():
        assert sha(Path(p))==v['sha256'] and oct(Path(p).stat().st_mode&0o777)==v['modo']
    manifest={'fecha':datetime.now(ZoneInfo('America/Santiago')).isoformat(),'proceso':2026,'oferta':2027,'etapas_comparadas':'Etapa 1 octubre V06 vs preparacion Etapa 2 noviembre V03','fuentes':inputs,'controles':control,'script':{'ruta':str(Path(__file__)),'sha256':sha(Path(__file__))},'salida':{'ruta':str(output),'sha256':sha(output)},'alcance':'Comparacion nominal y revision tematica orientativa. No acredita equivalencia curricular ni decision de cierre. Ninguna ausencia nominal demuestra por si sola que deba ofertarse o reincorporarse.'}
    write_once(ROOT/'11_gobernanza'/tag/'MANIFIESTO.json',(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n').encode())
    print(json.dumps({'controles':control,'salida':str(output)},ensure_ascii=False))


if __name__=='__main__':main()
