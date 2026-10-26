"""Compare names without editing SIES records or inferring program equivalence."""

import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import ROOT, tsv, write_once


def normalized(value):
    value = ''.join(c for c in unicodedata.normalize('NFD', value.upper()) if unicodedata.category(c) != 'Mn')
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', value).split())


def core(value):
    # Comparison key only: never written back to source names or PES columns.
    return ' '.join(t for t in normalized(value).split() if t not in {'DIPLOMADO', 'EN', 'DE', 'DEL', 'LA', 'EL', 'LOS', 'LAS', 'Y', 'PARA', 'CON'})


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    source = ROOT / '00_fuentes_originales/etapa1_octubre_20261020/reporte_dinamico_6047_2026-10-20_09:05:04.csv'
    xlsx = ROOT / '03_fuentes_institucionales/entregas_docencia/20261023_diplomados_vf/Oferta Diplomados 2027 - SIES vf.xlsx'
    base = ROOT / '07_resultados/borradores_prueba_no_oficiales/etapa1_octubre_20261020/PRUEBA_PES_OFERTA_2027_ETAPA1_17467_20261020_V04_SIN_SALTO_FINAL.csv'
    hashes = {str(p): sha(p) for p in (source, xlsx, base)}
    sies = [r for r in csv.DictReader(io.StringIO(source.read_bytes().decode('cp1252')), delimiter=';') if r['COD_NIVEL_GLOBAL'] == '3']
    ws = load_workbook(xlsx, data_only=False)['Hoja2']
    doc = [(i, ws.cell(i, 2).value) for i in range(5, ws.max_row + 1) if ws.cell(i, 2).value]
    results = []
    for r in sies:
        name = r['NOMBRE_CARRERA']
        exact = [(i, n) for i, n in doc if core(name) == core(n)]
        assert len(exact) <= 1
        expanded = [(i, n) for i, n in doc if set(core(name).split()) < set(core(n).split())] if not exact else []
        candidates = exact or expanded
        state = 'COINCIDENCIA_NOMINAL_NORMALIZADA' if exact else ('NOMBRE_AMPLIADO_REQUIERE_CONFIRMACION' if expanded else 'SIN_COINCIDENCIA_NOMINAL')
        results.append([r['COD_CARRERA'], name, normalized(name), core(name), state, ' | '.join(str(i) for i, _ in candidates), ' | '.join(n for _, n in candidates), r['VIGENCIA_CARRERA'], 'SIN_CAMBIOS'])
    counts = dict(Counter(r[4] for r in results))
    assert counts['COINCIDENCIA_NOMINAL_NORMALIZADA'] == 5
    assert len(results) == 21
    out = ROOT / '06_validaciones/etapa1_octubre_20261023_diplomados/conciliacion_normalizada_v02'
    gov = ROOT / '11_gobernanza/etapa1_octubre_20261023_diplomados/conciliacion_normalizada_v02'
    tsv(out / 'CONCILIACION_NOMBRES.tsv', [['cod_carrera', 'nombre_sies', 'normalizado', 'clave_comparacion', 'estado', 'filas_excel', 'nombres_excel', 'vigencia_original', 'accion']] + results)
    note = '''# Confirmaciones y conciliacion nominal

23-10-2026. Etapa 1 octubre, Oferta Academica 2027, proceso 2026.
Nueva instruccion del usuario: 15 vacantes en primer semestre y 15 en segundo.
Esto sustituye su indicacion previa de informar un solo total. Son 30 cupos por
programa, no los 15 totales que registra la fuente; es una decision interna
explicita, no una regla del manual ni una lectura del Excel.
El usuario solicita ver la conciliacion antes de cambiar vigencias.

Se normalizan mayusculas, tildes, espacios y puntuacion en ambos listados y se
comparan tambien claves sin articulos, preposiciones y el prefijo DIPLOMADO.
Los nombres originales no se alteran. Se obtienen 5 coincidencias nominales.
Cuando el nombre SIES esta contenido en un nombre mas largo se muestra como
candidato para confirmacion, no como equivalencia autorizada. La similitud
tematica o de caracteres no basta para transferir datos a un codigo SIES.

No se cambia ninguna vigencia ni se genera un nuevo CSV. La ausencia de nombre
en el Excel no equivale a vigencia 2 o 3. Los casos sin correspondencia siguen
pendientes de decision; no se distribuyen datos entre programas por parecido.
Para una futura carga se aplicara 15/15 a los programas cuya correspondencia
quede establecida, registrando separadamente la diferencia respecto del Excel.
Respaldo local; sin commit, push o envio a PES.
'''
    write_once(gov / 'DECISION_Y_METODO.md', note.encode('utf-8'))
    manifest = {'estado': 'CONCILIACION_PARA_REVISION_SIN_CAMBIOS_CARGA', 'entradas': hashes, 'script': {'ruta': str(Path(__file__).resolve()), 'sha256': sha(Path(__file__))}, 'controles': counts, 'vacantes_autorizadas_por_programa': {'primer_semestre': 15, 'segundo_semestre': 15}, 'cambios_vigencia': 0, 'salidas': [{'ruta': str(p), 'sha256': sha(p)} for p in (out / 'CONCILIACION_NOMBRES.tsv', gov / 'DECISION_Y_METODO.md')]}
    write_once(gov / 'MANIFIESTO.json', (json.dumps(manifest, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))
    assert all(sha(Path(p)) == h for p, h in hashes.items())
    print(json.dumps({'controles': counts, 'nombres_ampliados': [r for r in results if r[4] == 'NOMBRE_AMPLIADO_REQUIERE_CONFIRMACION']}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
