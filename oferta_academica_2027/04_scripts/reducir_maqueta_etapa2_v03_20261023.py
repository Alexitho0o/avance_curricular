"""Preserve source evidence, remove the unsupported worksheet via OOXML, and verify."""
import io
import json
import posixpath
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import ROOT, sha, write_once

TAG='maqueta_etapa2_noviembre_20261023_v03'
AUDIT=ROOT/'06_validaciones'/TAG
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def prepare():
    source=Path.home()/'Desktop/DIPLOMADOS_ETAPA2_PARA_COMPLETAR_20261023_V02.xlsx'
    wb=load_workbook(source,data_only=False)
    cached=load_workbook(source,data_only=True)
    assert wb.sheetnames==['Programas','Pendientes por campo','Diplomados para completar']
    resolved=[]
    for sheet in wb:
        if sheet.title=='Programas':continue
        for row in sheet:
            for cell in row:
                if cell.data_type=='f' and 'Programas' in cell.value:
                    # Only the simple references in this reviewed template can be materialized.
                    import re
                    m=re.fullmatch(r"='Programas'!([A-Z]+[0-9]+)",cell.value)
                    assert m,cell.value
                    value=wb['Programas'][m[1]].value
                    assert isinstance(value,(int,float)) and cached[sheet.title][cell.coordinate].value==value
                    resolved.append({'hoja':sheet.title,'celda':cell.coordinate,'formula_anterior':cell.value,'valor':value})
    assert len(resolved)==110
    frozen=ROOT/'07_resultados/cargas_congeladas/20261023_etapa1_octubre_v06_aceptada'
    inputs={str(p):{'sha256':sha(p),'modo':oct(p.stat().st_mode&0o777)} for p in [source]+sorted(p for p in frozen.iterdir() if p.is_file())}
    snapshot=ROOT/'00_fuentes_originales/etapa2_noviembre_20261023_maqueta_usuario'/source.name
    write_once(snapshot,source.read_bytes())
    payload={'source':str(source),'snapshot':str(snapshot),'fuentes':inputs,'referencias_resueltas':resolved}
    write_once(AUDIT/'INSUMOS.json',(json.dumps(payload,ensure_ascii=False,indent=2)+'\n').encode())
    print('110 referencias verificadas; fuente y congelado registrados.')


def finalize():
    data=json.loads((AUDIT/'INSUMOS.json').read_text())
    intermediate=AUDIT/'INTERMEDIO_TRES_HOJAS.xlsx'
    edits={};remove=set()
    with zipfile.ZipFile(intermediate) as z:
        book=ET.fromstring(z.read('xl/workbook.xml'))
        rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))
        content=ET.fromstring(z.read('[Content_Types].xml'))
        sheets=book.find(f'{{{NS}}}sheets')
        victim=next(s for s in sheets if s.get('name')=='Programas')
        rid=victim.get(f'{{{RNS}}}id')
        rel=next(r for r in rels if r.get('Id')==rid)
        target=rel.get('Target')
        part=target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        remove.add(part)
        relpart=posixpath.dirname(part)+'/_rels/'+posixpath.basename(part)+'.rels'
        if relpart in z.namelist():
            remove.add(relpart)
            for item in ET.fromstring(z.read(relpart)):
                if item.get('TargetMode')=='External':continue
                t=item.get('Target')
                remove.add(t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.dirname(part)+'/'+t))
        idx=list(sheets).index(victim)
        sheets.remove(victim);rels.remove(rel)
        for views in book.findall(f'{{{NS}}}bookViews/{{{NS}}}workbookView'):
            views.set('activeTab','0');views.set('firstSheet','0')
        names=book.find(f'{{{NS}}}definedNames')
        if names is not None:
            for n in list(names):
                local=n.get('localSheetId')
                if local is not None:
                    if int(local)==idx:names.remove(n)
                    elif int(local)>idx:n.set('localSheetId',str(int(local)-1))
                assert 'Programas' not in (n.text or '')
        for rel in list(rels):
            if rel.get('Type','').endswith('/calcChain'):
                t=rel.get('Target');remove.add(t.lstrip('/') if t.startswith('/') else 'xl/'+t);rels.remove(rel)
        for item in list(content):
            if item.get('PartName','').lstrip('/') in remove:content.remove(item)
        for part,root in [('xl/workbook.xml',book),('xl/_rels/workbook.xml.rels',rels),('[Content_Types].xml',content)]:
            edits[part]=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        # Restore native internal navigation, which the authoring engine can omit on import.
        targets={r.get('Id'):r.get('Target') for r in rels}
        for sh in sheets:
            t=targets[sh.get(f'{{{RNS}}}id')];part=t.lstrip('/') if t.startswith('/') else 'xl/'+t
            root=ET.fromstring(z.read(part))
            existing=root.find(f'{{{NS}}}hyperlinks')
            if existing is not None:root.remove(existing)
            links=ET.Element(f'{{{NS}}}hyperlinks')
            for i in range(48):
                c=chr(65+i) if i<26 else 'A'+chr(65+i-26);r=i+9
                refs=[f'B{r}',f'I{r}'] if sh.get('name')=='Pendientes por campo' else [f'{c}7',f'{c}8']
                location=f"'Diplomados para completar'!{c}8" if sh.get('name')=='Pendientes por campo' else f"'Pendientes por campo'!B{r}"
                for ref in refs:ET.SubElement(links,f'{{{NS}}}hyperlink',ref=ref,location=location)
            predecessors={'sheetPr','dimension','sheetViews','sheetFormatPr','cols','sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews','mergeCells','phoneticPr','conditionalFormatting','dataValidations'}
            at=next((i for i,e in enumerate(root) if ET.QName(e).localname not in predecessors),len(root));root.insert(at,links)
            for v in root.findall(f'{{{NS}}}sheetViews/{{{NS}}}sheetView'):v.set('tabSelected','1' if sh.get('name')=='Pendientes por campo' else '0')
            edits[part]=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
        if 'docProps/app.xml' in z.namelist():
            app=ET.fromstring(z.read('docProps/app.xml'))
            appns=ET.QName(app).namespace
            vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
            titles=app.find(f'{{{appns}}}TitlesOfParts/{{{vt}}}vector')
            if titles is not None:
                for e in list(titles):titles.remove(e)
                titles.set('size','2')
                for s in sheets:ET.SubElement(titles,f'{{{vt}}}lpstr').text=s.get('name')
            pairs=app.find(f'{{{appns}}}HeadingPairs/{{{vt}}}vector')
            if pairs is not None:
                for j,e in enumerate(pairs):
                    if ''.join(e.itertext())=='Worksheets' and j+1<len(pairs):
                        n=pairs[j+1].find(f'{{{vt}}}i4')
                        if n is not None:n.text='2'
            edits['docProps/app.xml']=ET.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as out:
            for info in z.infolist():
                if info.filename not in remove:out.writestr(info,edits.get(info.filename,z.read(info.filename)))
    name='DIPLOMADOS_ETAPA2_PARA_COMPLETAR_20261023_V03.xlsx'
    destinations=[ROOT/'10_entrega'/TAG/name,Path.home()/'Desktop'/name]
    for p in destinations:write_once(p,b.getvalue())
    old=load_workbook(data['source'],data_only=True)
    new=load_workbook(destinations[0],data_only=True)
    formulas=load_workbook(destinations[0],data_only=False)
    assert new.sheetnames==['Pendientes por campo','Diplomados para completar']
    for r in range(5,31):
        for c in range(1,49):assert old['Diplomados para completar'].cell(r,c).value==new['Diplomados para completar'].cell(r,c).value
    links=0
    for s in formulas:
        for row in s:
            for c in row:
                if c.data_type=='f':assert '#REF!' not in c.value and 'Programas' not in c.value
                if c.hyperlink:
                    assert 'Programas' not in c.hyperlink.location;links+=1
    assert links==192
    assert not [(s.title,c.coordinate) for s in new for row in s for c in row if c.data_type=='e']
    for p,info in data['fuentes'].items():
        assert sha(Path(p))==info['sha256'] and oct(Path(p).stat().st_mode&0o777)==info['modo']
    report={'estado':'PLANTILLA_DOS_HOJAS_PARA_COMPLETAR','fuentes':data['fuentes'],'hojas':new.sheetnames,'datos_22_diplomados_conservados':True,'referencias_materializadas':data['referencias_resueltas'],'enlaces_nativos':links,'errores_formula':0,'etapa1_intacta':True,'salidas':[{'ruta':str(p),'sha256':sha(p)} for p in destinations],'criterio_vigencia3':'No reincorporar automaticamente. Manual p.29 excluye no vigentes de oferta validada; Etapa 2 informa nuevos/faltantes. Confirmar oferta efectiva y correspondencias, sin asumir reactivacion ni equivalencias.', 'scripts':[{'ruta':str(Path(__file__)),'sha256':sha(Path(__file__))},{'ruta':str(ROOT/'04_scripts/excel_manual_diplomados_20261023/maqueta_etapa2_v03.mjs'),'sha256':sha(ROOT/'04_scripts/excel_manual_diplomados_20261023/maqueta_etapa2_v03.mjs')}]}
    write_once(ROOT/'11_gobernanza'/TAG/'MANIFIESTO.json',(json.dumps(report,ensure_ascii=False,indent=2)+'\n').encode())
    print(json.dumps({'hojas':new.sheetnames,'datos_conservados':True,'referencias_rotas':0,'enlaces':links,'salidas':[str(p) for p in destinations]},ensure_ascii=False))


if __name__=='__main__':
    prepare() if sys.argv[1]=='prepare' else finalize()
