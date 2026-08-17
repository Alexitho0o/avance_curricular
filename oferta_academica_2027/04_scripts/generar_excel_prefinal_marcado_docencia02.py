#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
PREFINAL = ROOT / "05_datos_trabajo" / "PREFINAL_ETAPA1_DOCENCIA02_48_COLUMNAS_CON_ENCABEZADO_20260814_081424_v2.csv"
VALIDACIONES_DIR = ROOT / "06_validaciones"
OUT_ENTREGA = ROOT / "10_entrega"
OUT_DESKTOP = Path.home() / "Desktop"

AREAS = [
    "AREA_ADMIN_DERECHO",
    "AREA_AGRI_SILVI_PESCA_VET",
    "AREA_ARTES_HUMANIDADES",
    "AREA_CIENCIAS_NAT_MAT_ESTAD",
    "AREA_CS_SOCIAL_PERIODISMO_INFO",
    "AREA_EDUCACION",
    "AREA_INGE_INDUSTRIA_CONSTRUC",
    "AREA_SALUD_BIENESTAR",
    "AREA_SERVICIOS",
    "AREA_TECNO_INFO_COMUNICA",
]


def latest_validation() -> tuple[Path, Path]:
    jsons = sorted(VALIDACIONES_DIR.glob("VALIDACION_REGLAS_MANUAL_ETAPA1_DOCENCIA02_*.json"))
    if not jsons:
        raise FileNotFoundError("No existe validacion de reglas manual.")
    json_path = jsons[-1]
    tsv_path = json_path.with_suffix(".tsv")
    if not tsv_path.exists():
        raise FileNotFoundError(f"No existe TSV asociado: {tsv_path}")
    return json_path, tsv_path


def read_prefinal() -> tuple[list[str], list[list[str]]]:
    with PREFINAL.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    return rows[0], rows[1:]


def read_issues(tsv_path: Path) -> list[dict[str, str]]:
    with tsv_path.open("r", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def issue_fields(field: str, headers: list[str]) -> list[str]:
    if field == "areas_destino":
        return [f for f in AREAS if f in headers]
    fields = [f for f in field.split(",") if f]
    return [f for f in fields if f in headers]


def autosize(ws, max_width: int = 42) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        width = 10
        for row_idx in range(1, min(ws.max_row, 120) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                width = max(width, min(max_width, len(str(value)) + 2))
        ws.column_dimensions[letter].width = width


def main() -> None:
    validation_json, validation_tsv = latest_validation()
    validation = json.loads(validation_json.read_text(encoding="utf-8"))
    summary = validation["resumen"]
    issues = read_issues(validation_tsv)
    headers, data_rows = read_prefinal()

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_name = f"OFERTA_ACADEMICA_ETAPA1_PREFINAL_MARCADO_PENDIENTES_DOCENCIA02_{ts}.xlsx"
    OUT_ENTREGA.mkdir(parents=True, exist_ok=True)
    OUT_DESKTOP.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_data = wb.create_sheet("Archivo 48 columnas")
    ws_issues = wb.create_sheet("Hallazgos")

    red = PatternFill("solid", fgColor="FFC7CE")
    red_dark = PatternFill("solid", fgColor="E06666")
    amber = PatternFill("solid", fgColor="FCE4D6")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")
    gray_fill = PatternFill("solid", fgColor="E7E6E6")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(bottom=thin_gray)

    # Resumen
    ws_resumen["A1"] = "Oferta Académica 2027 - Etapa 1 prefinal marcado"
    ws_resumen["A1"].font = Font(bold=True, size=16, color="1F4E78")
    ws_resumen["A3"] = "Estado"
    ws_resumen["B3"] = "NO SUBIR A SIES: existen bloqueos pendientes"
    ws_resumen["B3"].fill = red
    ws_resumen["B3"].font = Font(bold=True, color="9C0006")

    rows_summary = [
        ("Registros", summary.get("total_registros")),
        ("Columnas estructura carga", summary.get("total_columnas")),
        ("Hallazgos totales", summary.get("issues_total")),
        ("Bloqueantes", summary.get("por_severidad", {}).get("BLOQUEANTE", 0)),
        ("Revisión", summary.get("por_severidad", {}).get("REVISION", 0)),
        ("Fuente prefinal", str(PREFINAL)),
        ("Validación usada", str(validation_json)),
    ]
    for idx, (k, v) in enumerate(rows_summary, start=5):
        ws_resumen.cell(idx, 1).value = k
        ws_resumen.cell(idx, 2).value = v
        ws_resumen.cell(idx, 1).font = bold

    ws_resumen["A14"] = "Leyenda"
    ws_resumen["A14"].font = bold
    ws_resumen["A15"] = "Rojo"
    ws_resumen["B15"] = "Campo bloqueante: debe corregirse antes de generar CSV final para PES/SIES."
    ws_resumen["A15"].fill = red
    ws_resumen["A16"] = "Naranjo"
    ws_resumen["B16"] = "Campo en revisión: requiere confirmación Docencia/RAP/vacantes."
    ws_resumen["A16"].fill = amber
    ws_resumen["A17"] = "Gris"
    ws_resumen["B17"] = "Campos de estructura o referencia; no editar salvo instrucción expresa."
    ws_resumen["A17"].fill = gray_fill

    by_rule = Counter(i["regla_id"] for i in issues)
    ws_resumen["A20"] = "Resumen por regla"
    ws_resumen["A20"].font = bold
    ws_resumen.append([])
    start_rule = 21
    ws_resumen.cell(start_rule, 1).value = "Regla"
    ws_resumen.cell(start_rule, 2).value = "Cantidad"
    ws_resumen.cell(start_rule, 3).value = "Tipo"
    for c in range(1, 4):
        ws_resumen.cell(start_rule, c).fill = header_fill
        ws_resumen.cell(start_rule, c).font = white_font
    rule_types = {i["regla_id"]: i["severidad"] for i in issues}
    for r, (rule, count) in enumerate(sorted(by_rule.items()), start=start_rule + 1):
        ws_resumen.cell(r, 1).value = rule
        ws_resumen.cell(r, 2).value = count
        ws_resumen.cell(r, 3).value = rule_types.get(rule, "")

    # Data sheet
    ws_data.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws_data.cell(1, c)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in data_rows:
        ws_data.append(row)

    header_index = {h: i + 1 for i, h in enumerate(headers)}
    comments_by_cell: dict[tuple[int, int], list[str]] = defaultdict(list)

    for issue in issues:
        fila = issue.get("fila", "")
        if not fila.isdigit():
            continue
        excel_row = int(fila)
        fields = issue_fields(issue.get("campo", ""), headers)
        if not fields:
            continue
        for field in fields:
            col = header_index[field]
            cell = ws_data.cell(excel_row, col)
            if issue["severidad"] == "BLOQUEANTE":
                cell.fill = red
                cell.font = Font(color="9C0006")
            else:
                cell.fill = amber
                cell.font = Font(color="7F6000")
            comments_by_cell[(excel_row, col)].append(f"{issue['regla_id']} {issue['severidad']}: {issue['detalle']}")

    for (row, col), comments in comments_by_cell.items():
        ws_data.cell(row, col).comment = Comment("\n".join(comments[:8]), "Codex")

    for row in ws_data.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")
    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = ws_data.dimensions

    # Hallazgos
    issue_headers = [
        "regla_id",
        "severidad",
        "fila",
        "cod_carrera",
        "nombre_carrera",
        "modalidad",
        "cod_jornada",
        "version",
        "campo",
        "valor",
        "detalle",
    ]
    ws_issues.append(issue_headers)
    for c in range(1, len(issue_headers) + 1):
        cell = ws_issues.cell(1, c)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for issue in issues:
        ws_issues.append([issue.get(h, "") for h in issue_headers])
        fill = red if issue.get("severidad") == "BLOQUEANTE" else amber
        for c in range(1, len(issue_headers) + 1):
            ws_issues.cell(ws_issues.max_row, c).fill = fill
            ws_issues.cell(ws_issues.max_row, c).border = border
            ws_issues.cell(ws_issues.max_row, c).alignment = Alignment(vertical="top", wrap_text=True)
    ws_issues.freeze_panes = "A2"
    ws_issues.auto_filter.ref = ws_issues.dimensions

    for ws in [ws_resumen, ws_data, ws_issues]:
        ws.sheet_view.showGridLines = False
        autosize(ws)
    ws_data.column_dimensions["D"].width = 38
    ws_issues.column_dimensions["E"].width = 38
    ws_issues.column_dimensions["K"].width = 68
    ws_resumen.column_dimensions["A"].width = 28
    ws_resumen.column_dimensions["B"].width = 95

    desktop_path = OUT_DESKTOP / output_name
    entrega_path = OUT_ENTREGA / output_name
    wb.save(desktop_path)
    wb.save(entrega_path)

    print(json.dumps({
        "desktop": str(desktop_path),
        "entrega": str(entrega_path),
        "validacion": str(validation_json),
        "hallazgos": len(issues),
        "bloqueantes": summary.get("por_severidad", {}).get("BLOQUEANTE", 0),
        "revision": summary.get("por_severidad", {}).get("REVISION", 0),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
