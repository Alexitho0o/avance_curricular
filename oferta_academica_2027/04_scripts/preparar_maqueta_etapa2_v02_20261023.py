"""Snapshot the user's revised layout and prepare a diploma-only completion grid."""

import json
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from preparar_etapa1_octubre_20261020 import ROOT, FIELDS, sha, write_once


def main():
    source = Path.home() / 'Desktop/REVISION_DIPLOMADOS_ETAPA2_NOVIEMBRE_20261023.xlsx'
    snapshot = ROOT / '00_fuentes_originales/etapa2_noviembre_20261023_maqueta_usuario' / source.name
    manual = ROOT / '00_fuentes_originales/etapa1_octubre_20261020/Instructivo_Oferta_Academica_2027_CFT_E_IP_DFE.pdf'
    frozen = ROOT / '07_resultados/cargas_congeladas/20261023_etapa1_octubre_v06_aceptada'
    inputs = {str(p): {'sha256': sha(p), 'modo': oct(p.stat().st_mode & 0o777)} for p in [source, manual] + sorted(p for p in frozen.iterdir() if p.is_file())}
    write_once(snapshot, source.read_bytes())
    wb = load_workbook(source, data_only=False)
    assert wb.sheetnames == ['Programas', 'Pendientes por campo']
    ws = wb['Programas']
    expected = ['Programa Docencia','Tratamiento','Codigo SIES V06','Programa SIES V06','Observacion','Periodo','Vacantes 1er y 2do semestre','Col. D: Matricula original','Titulacion original','Certificado original','Arancel original','Web']
    assert [ws.cell(8, c).value for c in range(1, 13)] == expected
    old = json.loads((ROOT / '06_validaciones/preparacion_etapa2_noviembre_20261023/ANALISIS.json').read_text())
    candidates = []
    issues = {
        13: 'Confirmar escritura de "Analytics y business inteligence"; no se corrige el nombre por inferencia.',
        20: 'Confirmar nombre completo sin siglas UX/UI. Solo se retiro la barra en la preparacion de formato.',
        27: 'Confirmar nombre completo sin abreviatura TI; no se expande por inferencia.',
        29: 'Confirmar denominacion: la fuente repite "Diplomado en Diplomado en". No se elimino una parte del nombre.',
        31: 'Se retiro puntuacion final y espacio no separable para ajustar el formato SIES.',
    }
    for r in range(9, 38):
        original = old['programas'][r-9]
        assert ws.cell(r,1).value == original[0]
        assert ws.cell(r,2).value == original[1]
        if ws.cell(r,2).value != 'CANDIDATO ETAPA 2':
            continue
        name = ws.cell(r,1).value
        stripped = ''.join(c for c in unicodedata.normalize('NFD', name.upper()) if not unicodedata.combining(c))
        normalized = ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', stripped).split())
        assert ws.cell(r,6).value == 2027
        assert ws.cell(r,7).value == 15
        assert ws.cell(r,8).value in [400000,600000,990000]
        assert all(ws.cell(r,c).value == 0 for c in [9,10,11])
        candidates.append({'fila_maqueta':r,'fila_docencia':original[5],'nombre_original':name,'nombre_sies_preparado':normalized,'observacion':issues.get(original[5],''),'importe':ws.cell(r,8).value})
    assert len(candidates) == 22
    pages = PdfReader(manual).pages
    schema = []
    for p in range(35,42):
        schema.extend(re.findall(r'^\s*(?:[A-Z]|A[A-V])\s+([A-Z][A-Z_]+)\b',pages[p-1].extract_text(),re.MULTILINE))
    assert schema == FIELDS
    preserved = {s.title: {c.coordinate:c.value for row in s for c in row if c.value is not None} for s in wb}
    data = {'source':str(source),'snapshot':str(snapshot),'inputs':inputs,'fields':FIELDS,'candidatos':candidates,'reglas':old['reglas'],'contenido_original':preserved,
            'blank_indices':[2]+list(range(34,42)),
            'fuente_normativa':str(manual),'estado':'PLANTILLA_PARA_COMPLETAR_NO_CARGA',
            'decisiones':{'niveles':'Usuario confirma que todos son diplomados. 3/5 precargados; verificar duracion desde un semestre, manual p.37.',
                          'vacantes':'Maqueta del usuario columna G: Vacantes 1er y 2do semestre=15; coincide con instruccion previa 15+15.',
                          'cobros':'Se aplica criterio anterior del usuario: monto etiquetado Matricula corresponde a arancel y matricula=0. Marcado para revisar, no aceptacion de nuevos programas.',
                          'vigencia':'Se deja en blanco para confirmar oferta de cada programa; no cambia ninguna vigencia de Etapa 1.',
                          'anio':'Periodo 2027 de la maqueta no se traslada a ANIO_INICIO.'}}
    out = ROOT / '06_validaciones/maqueta_etapa2_noviembre_20261023_v02/INSUMOS.json'
    write_once(out,(json.dumps(data,ensure_ascii=False,indent=2)+'\n').encode())
    print(json.dumps({'candidatos':len(candidates),'campos':len(FIELDS),'obligatorios_con_dato':39,'vacios_exigidos':9,'fuente_usuario_resguardada':str(snapshot)},ensure_ascii=False))


if __name__ == '__main__':
    main()
