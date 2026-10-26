"""Prepare accepted vigencia-3 rows and verify an added review sheet."""
import csv
import io
import json
import posixpath
import sys
import zipfile
from pathlib import Path
from lxml import etree as ET
from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import ROOT, FIELDS, sha, write_once

TAG='maqueta_etapa2_noviembre_20261023_v04'
AUDIT=ROOT/'06_validaciones'/TAG
NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS='http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def prepare():
    source=Path.home()/'Desktop/DIPLOMADOS_ETAPA2_PARA_COMPLETAR_20261023_V03.xlsx'
    frozen=ROOT/'07_resultados/cargas_congeladas/20261023_etapa1_octubre_v06_aceptada'
    base=frozen/'CARGA_OFERTA_2027_ETAPA1_17467_20261023_V06.csv'
    assert sha(base)=='d7fe88035cd1215d448272a9854bee0932c45c200f4f80bb58636b62a6ae4e0b'
    inputs={str(p):{'sha256':sha(p),'modo':oct(p.stat().st_mode&0o777)} for p in [source]+sorted(p for p in frozen.iterdir() if p.is_file())}
    wb=load_workbook(source,data_only=False)
    assert wb.sheetnames==['Pendientes por campo','Diplomados para completar']
    rows=list(csv.reader(io.StringIO(base.read_text()),delimiter=';'))
    assert len(rows)==161 and all(len(r)==48 for r in rows)
    records=[dict(zip(FIELDS,r)) for r in rows]
    lows=[r for r in records if r['VIGENCIA_CARRERA']=='3' and r['COD_NIVEL_GLOBAL']=='3' and r['COD_NIVEL_CARRERA']=='5']
    assert {r['COD_CARRERA'] for r in lows}=={'40','63','64','65','67','71','95','97','98','101','102','103','104','105'}
    data={'source':str(source),'fuentes':inputs,'rows':[[r['COD_CARRERA'],r['NOMBRE_CARRERA'],3,None,None] for r in sorted(lows,key=lambda r:int(r['COD_CARRERA']))]}
    write_once(AUDIT/'INSUMOS.json',(json.dumps(data,ensure_ascii=False,indent=2)+'\n').encode())
    print('14 filas verificadas contra V06 aceptada; ninguna vigencia 1 o 2.')


def sheetpaths(z):
    book=ET.fromstring(z.read('xl/workbook.xml'))
    rels={r.get('Id'):r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    return {s.get('name'):(rels[s.get(f'{{{RNS}}}id')].lstrip('/') if rels[s.get(f'{{{RNS}}}id')].startswith('/') else posixpath.normpath('xl/'+rels[s.get(f'{{{RNS}}}id')])) for s in book.find(f'{{{NS}}}sheets')}


def finalize():
    data=json.loads((AUDIT/'INSUMOS.json').read_text())
    intermediate=AUDIT/'INTERMEDIO_V04.xlsx'
    # Preserve native internal links unsupported by the workbook authoring engine.
    edits={}
    with zipfile.ZipFile(data['source']) as old, zipfile.ZipFile(intermediate) as new:
        op,np=sheetpaths(old),sheetpaths(new)
        for name,part in op.items():
            source=ET.fromstring(old.read(part));dest=ET.fromstring(new.read(np[name]))
            oldlinks=source.find(f'{{{NS}}}hyperlinks')
            existing=dest.find(f'{{{NS}}}hyperlinks')
            if existing is not None:dest.remove(existing)
            if oldlinks is not None:
                before={'sheetPr','dimension','sheetViews','sheetFormatPr','cols','sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews','mergeCells','phoneticPr','conditionalFormatting','dataValidations'}
                at=next((i for i,e in enumerate(dest) if ET.QName(e).localname not in before),len(dest))
                dest.insert(at,oldlinks)
            edits[np[name]]=ET.tostring(dest,xml_declaration=True,encoding='UTF-8',standalone=True)
        buffer=io.BytesIO()
        with zipfile.ZipFile(buffer,'w') as out:
            for part in new.infolist():out.writestr(part,edits.get(part.filename,new.read(part.filename)))
    name='DIPLOMADOS_ETAPA2_PARA_COMPLETAR_20261023_V04.xlsx'
    outputs=[ROOT/'10_entrega'/TAG/name,Path.home()/'Desktop'/name]
    for p in outputs:write_once(p,buffer.getvalue())
    a=load_workbook(data['source'],data_only=False);b=load_workbook(outputs[0],data_only=False);cached=load_workbook(outputs[0],data_only=True)
    assert b.sheetnames==a.sheetnames+['Diplomados vigencia 3']
    unchanged=0
    for s in a:
        for row in s:
            for c in row:
                x=b[s.title][c.coordinate]
                assert c.value==x.value,(s.title,c.coordinate)
                assert (c.font.name,c.font.sz,c.font.bold,c.number_format,c.fill.fgColor.rgb,c.alignment.horizontal,c.alignment.wrap_text)==(x.font.name,x.font.sz,x.font.bold,x.number_format,x.fill.fgColor.rgb,x.alignment.horizontal,x.alignment.wrap_text)
                assert (c.hyperlink.location if c.hyperlink else None)==(x.hyperlink.location if x.hyperlink else None)
                unchanged+=1
        assert s.freeze_panes==b[s.title].freeze_panes
    review=b['Diplomados vigencia 3']
    assert review.max_row==22 and review.max_column==5
    for n,values in enumerate(data['rows'],9):assert [review.cell(n,c).value for c in range(1,6)]==values
    assert len(review.data_validations.dataValidation)==1
    assert next(iter(review.data_validations.dataValidation)).formula1=='"SI,NO"'
    assert cached[review.title]['A5'].value=='Pendientes de respuesta: 14. Solicitan revision: 0. Sin solicitud: 0.'
    assert not [(s.title,c.coordinate,c.value) for s in cached for row in s for c in row if c.data_type=='e']
    for p,x in data['fuentes'].items():assert sha(Path(p))==x['sha256'] and oct(Path(p).stat().st_mode&0o777)==x['modo']
    report={'estado':'REVISION_DOCENCIA_NO_CARGA','fuentes':data['fuentes'],'filas_vigencia3':14,'cambios_a_las_dos_hojas_previas':0,'celdas_previas_verificadas':unchanged,'etapa1_congelada_intacta':True,'decision_rectificacion':'En blanco; SI o NO lo responde Docencia. No modifica PES ni garantiza admision de rectificacion.','formulas_sin_errores':True,'salidas':[{'ruta':str(p),'sha256':sha(p)} for p in outputs],'scripts':[{'ruta':str(Path(__file__)),'sha256':sha(Path(__file__))},{'ruta':str(ROOT/'04_scripts/excel_manual_diplomados_20261023/maqueta_etapa2_v04.mjs'),'sha256':sha(ROOT/'04_scripts/excel_manual_diplomados_20261023/maqueta_etapa2_v04.mjs')}]}
    write_once(ROOT/'11_gobernanza'/TAG/'MANIFIESTO.json',(json.dumps(report,ensure_ascii=False,indent=2)+'\n').encode())
    print(json.dumps({'filas':14,'hojas':b.sheetnames,'cambios_previas':0,'salidas':[str(p) for p in outputs]},ensure_ascii=False))


if __name__=='__main__':prepare() if sys.argv[1]=='prepare' else finalize()
