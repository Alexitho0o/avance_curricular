"""Apply user-approved diploma dispositions and institutional amounts to V04."""

import csv
import io
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader
from preparar_etapa1_octubre_20261020 import FIELDS, FINANCE, KEY, ROOT, sha, tsv, write_once


def main():
    base = ROOT / '07_resultados/borradores_prueba_no_oficiales/etapa1_octubre_20261020/PRUEBA_PES_OFERTA_2027_ETAPA1_17467_20261020_V04_SIN_SALTO_FINAL.csv'
    source = ROOT / '00_fuentes_originales/etapa1_octubre_20261020/reporte_dinamico_6047_2026-10-20_09:05:04.csv'
    manual = source.parent / 'Instructivo_Oferta_Academica_2027_CFT_E_IP_DFE.pdf'
    doc = ROOT / '03_fuentes_institucionales/entregas_docencia/20261023_diplomados_vf/Oferta Diplomados 2027 - SIES vf.xlsx'
    inputs = {str(p): sha(p) for p in (base, source, manual, doc)}
    before = list(csv.reader(io.StringIO(base.read_text(encoding='utf-8')), delimiter=';', strict=True))
    assert len(before) == 161 and all(len(r) == 48 for r in before)
    raw = list(csv.DictReader(io.StringIO(source.read_bytes().decode('cp1252')), delimiter=';'))
    rows = [dict(zip(FIELDS, r)) for r in before]
    assert len(raw) == len(rows)
    for a, b in zip(raw, rows):
        assert all(a[c] == b[c] for c in KEY)
    pages = PdfReader(manual).pages
    schema = []
    for n in range(21, 29):
        schema.extend(re.findall(r'^\s*(?:[A-Z]|A[A-V])\s+([A-Z][A-Z_]+)\b', pages[n-1].extract_text(), re.MULTILINE))
    assert schema == FIELDS

    inactive = {'40','63','64','65','67','71','95','97','98','101','102','103','104','105'}
    mapping = {
        '53': (17, 'Diplomado en Ciberseguridad Aplicada', 'Coincidencia nominal'),
        '66': (21, 'Diplomado en Fullstack', 'Coincidencia nominal'),
        '68': (19, 'Diplomado en Seguridad Ofensiva y ethical hacking', 'Equivalencia confirmada por usuario; no renombrar SIES'),
        '69': (18, 'Diplomado en Infraestructura Cloud', 'Coincidencia nominal normalizada'),
        '70': (25, 'Diplomado en Redes Industriales', 'Coincidencia nominal normalizada'),
        '96': (32, 'Diplomado en Gobernanza de Datos', 'Coincidencia nominal; no usar variante Gestion Estrategica'),
        '99': (9, 'Diplomado en Inteligencia Artificial Aplicada para la gestión publica y local', 'Candidato presentado en cuadro; usuario confirma restantes coincidentes salvo 95 y luego especifica bajas'),
    }
    diplomas = [r for r in rows if (r['COD_NIVEL_GLOBAL'], r['COD_NIVEL_CARRERA']) == ('3','5')]
    assert len(diplomas) == 21 and {r['COD_CARRERA'] for r in diplomas} == inactive | set(mapping)
    assert not inactive & set(mapping)
    ws = load_workbook(doc, data_only=False)['Hoja2']
    changes = []
    decisions = []
    protections = []
    for i, r in enumerate(rows):
        if r['COD_NIVEL_GLOBAL'] != '3':
            continue
        code = r['COD_CARRERA']
        assert r['COD_SEDE'] == '2' and r['VERSION'] == '1' and r['VIGENCIA_CARRERA'] == '1'
        if code in inactive:
            for field in ['CANTIDAD_MATRICULA_DFE','CANTIDAD_BENEFICIO_DFE']:
                value = raw[i][field]
                if value:
                    assert value.isdigit() and int(value) == 0, f'Proteccion positiva: {code} {field}'
                protections.append([code, field, value, 'Vacio: no acredita ausencia de proteccion' if not value else 'Sin cantidad positiva informada'])
            r['VIGENCIA_CARRERA'] = '3'
            reason = 'Decision explicita del usuario: vigencia 3; codigo 95 confirmado en mensaje posterior'
            decisions.append([code,r['NOMBRE_CARRERA'],'3','','',reason])
        else:
            excel_row, expected_name, reason = mapping[code]
            assert ws.cell(excel_row,2).value == expected_name
            assert r['FORMATO_VALOR'] == '1'
            for field, col in zip(FINANCE, [4,5,6,7]):
                cell = ws.cell(excel_row,col)
                assert cell.data_type != 'f' and isinstance(cell.value, int) and not isinstance(cell.value, bool) and cell.value >= 0
                r[field] = str(cell.value)
            r['VACANTES_PRIMER_SEMESTRE'] = '15'
            r['VACANTES_SEGUNDO_SEMESTRE'] = '15'
            assert ws.cell(excel_row,3).value == 15
            decisions.append([code,r['NOMBRE_CARRERA'],'1',excel_row,expected_name,reason])
        for j, field in enumerate(FIELDS):
            if before[i][j] != r[field]:
                trace = reason if field == 'VIGENCIA_CARRERA' else 'Usuario: 15 por semestre; no derivado del periodo Excel' if field.startswith('VACANTES_') else f'Docencia Hoja2 fila {mapping[code][0]}, columna {dict(zip(FINANCE,["D","E","F","G"]))[field]}'
                changes.append([i+1,code,r['NOMBRE_CARRERA'],field,before[i][j],r[field],trace])

    after = [[r[c] for c in FIELDS] for r in rows]
    assert len(changes) == 56
    assert sum(a != b for a, b in zip(before, after)) == 21
    assert all(a == b for a, b in zip(before, after) if a[FIELDS.index('COD_NIVEL_GLOBAL')] == '1')
    allowed = set(FINANCE + ['VACANTES_PRIMER_SEMESTRE','VACANTES_SEGUNDO_SEMESTRE','VIGENCIA_CARRERA'])
    assert all(c[3] in allowed for c in changes)
    assert len({tuple(r[k] for k in KEY) for r in rows}) == 161
    combos = {('1','1'),('1','2'),('1','5'),('2','1'),('2','2'),('2','3'),('2','5'),('3','4')}
    for r in rows:
        assert (r['MODALIDAD'],r['COD_JORNADA']) in combos
        assert (r['COD_NIVEL_GLOBAL'],r['COD_NIVEL_CARRERA']) in {('1','1'),('1','2'),('3','5'),('3','6')}
        assert r['VIGENCIA_CARRERA'] in {'1','2','3'}
        de, dt, total = [int(r[c]) for c in ['DURACION_ESTUDIOS','DURACION_TITULACION','DURACION_TOTAL']]
        assert de <= total <= de + dt and de <= 14
        if r['FECHA_ADMISION_INICIAL']:
            assert datetime.strptime(r['FECHA_ADMISION_INICIAL'],'%d/%m/%Y').strftime('%d/%m/%Y') == r['FECHA_ADMISION_INICIAL']
        if r['VIGENCIA_CARRERA'] == '1':
            assert int(r['VACANTES_PRIMER_SEMESTRE']) + int(r['VACANTES_SEGUNDO_SEMESTRE']) > 0
            assert all(r[c] != '' for c in FINANCE)
        if r['COD_NIVEL_GLOBAL'] == '3':
            assert all(r[c] == '' for c in FIELDS[34:42])
            if r['COD_CARRERA'] in mapping:
                assert r['VACANTES_PRIMER_SEMESTRE'] == r['VACANTES_SEGUNDO_SEMESTRE'] == '15'
            else:
                assert r['VIGENCIA_CARRERA'] == '3' and all(r[c] == '' for c in FINANCE)
                assert r['VACANTES_PRIMER_SEMESTRE'] == r['VACANTES_SEGUNDO_SEMESTRE'] == '0'
    def vacancies(group):
        return sum(int(r[c]) for r in group for c in ['VACANTES_PRIMER_SEMESTRE','VACANTES_SEGUNDO_SEMESTRE'])
    assert vacancies(rows) == 5671
    assert vacancies([r for r in rows if r['COD_NIVEL_GLOBAL']=='1']) == 5461
    assert vacancies(diplomas) == 210
    stream=io.StringIO(newline='')
    csv.writer(stream,delimiter=';',lineterminator='\n').writerows(after)
    payload=stream.getvalue().encode('utf-8')[:-1]
    assert not payload.endswith((b'\n',b'\r')) and not payload.startswith(b'\xef\xbb\xbf')
    assert list(csv.reader(io.StringIO(payload.decode('utf-8')),delimiter=';',strict=True)) == after
    assert len(payload.split(b'\n')) == 161
    name='CARGA_OFERTA_2027_ETAPA1_17467_20261023_V05.csv'
    destinations=[ROOT/'10_entrega/etapa1_octubre_20261023_v05'/name,Path.home()/'Desktop'/name]
    audit=ROOT/'06_validaciones/etapa1_octubre_20261023_v05'
    gov=ROOT/'11_gobernanza/etapa1_octubre_20261023_v05'
    for p in destinations:
        write_once(p,payload)
    tsv(audit/'CAMBIOS_POR_CELDA.tsv',[['fila_csv','cod_carrera','nombre_sies','campo','anterior','nuevo','origen_decision']]+changes)
    tsv(audit/'DECISIONES_21_DIPLOMADOS.tsv',[['cod_carrera','nombre_sies','vigencia','fila_excel','nombre_docencia','respaldo']]+decisions)
    tsv(audit/'REVISION_PROTECCIONES.tsv',[['cod_carrera','campo','valor_precarga','alcance']]+protections)
    controls={'filas':161,'columnas':48,'pregrado_sin_cambios':140,'diplomados_vigencia1':7,'diplomados_vigencia3':14,'vigencias':dict(Counter(r['VIGENCIA_CARRERA'] for r in rows)),'cambios_celdas':56,'vacantes_pregrado':5461,'vacantes_diplomados':210,'vacantes_total':5671,'vacios_financieros_solo_vigencia3':56,'duplicados_clave':0,'filas_vacias':0,'salto_final':False,'errores_pes_previos_vigencia1_reproducidos':0}
    write_once(audit/'CONTROLES.json',(json.dumps(controls,ensure_ascii=False,indent=2)+'\n').encode('utf-8'))
    note='''# Version V05 para carga PES, pendiente de aceptacion

23-10-2026. Oferta Academica 2027, proceso 2026, Etapa 1 octubre, ID 17467.
Base V04, no el Excel de edicion manual que fue entregado en blanco.
Vigencia 3 autorizada: 40,63,64,65,67,71,95,97,98,101,102,103,104,105.
La decision final sobre 40 sustituye su correspondencia previamente propuesta.
Vigencia 1 conservada: 53,66,68,69,70,96,99. No se cambian nombres ni claves.
Cinco cruces nominales y dos por confirmacion conversacional (68 y 99).
Para 99 se usa la fila 9 de Docencia: IA Aplicada para Gestion Publica y Local,
matricula $400.000. Los otros seis vigentes tienen matricula $990.000.
Titulacion, certificado y arancel: 0, tal como los entrega Docencia.
No se intercambia matricula con arancel y no se incorporan enlaces de diplomados.

Vacantes: 15 primer semestre y 15 segundo por instruccion explicita del usuario;
son 30 anuales por programa, aunque el Excel fuente informa un solo total de 15.
No se dedujo equivalencia entre periodos Excel y semestres SIES.
Los 14 no vigentes conservan 0 vacantes y los campos financieros vacios de SIES.
No se imputa 0 o -1 a esos importes. Los campos DFE vacios no prueban ausencia
de proteccion; PES confirmara la posibilidad de baja.

Validado: 161x48, orden manual pp.21-28, claves unicas, fechas DD/MM/AAAA,
pregrado integro, combinaciones modalidad/jornada y niveles, duraciones,
vacantes positivas y financieros completos para vigencia 1; sin linea final.
CSV UTF-8 sin BOM, punto y coma, sin encabezados, ninguna fila agregada/borrada.
Las condiciones observadas en los errores previos ya no aparecen para vigencia 1.
Esto no garantiza ausencia de otras validaciones de PES ni su aceptacion final.

Usar V05 para cargar; no esta aceptada ni congelada. Conservar el resultado de
PES y conciliar contra reporte 6048 antes de cerrar. No se envio a PES.
Respaldo local reproducible, sin commit ni push. Originales y V04 intactos.
'''
    write_once(gov/'ESTADO_Y_REANUDACION.md',note.encode('utf-8'))
    manifest={'fecha':datetime.now(ZoneInfo('America/Santiago')).isoformat(),'estado':'PREPARADO_PARA_CARGA_PENDIENTE_ACEPTACION_PES','entradas':[{'ruta':p,'sha256':h} for p,h in inputs.items()],'script':{'ruta':str(Path(__file__).resolve()),'sha256':sha(Path(__file__))},'controles':controls,'vigencia3':sorted(inactive,key=int),'correspondencias':mapping,'salidas':[{'ruta':str(p),'sha256':sha(p),'bytes':p.stat().st_size} for p in destinations]}
    if not (gov/'MANIFIESTO.json').exists():
        write_once(gov/'MANIFIESTO.json',(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n').encode('utf-8'))
    assert all(sha(Path(p))==h for p,h in inputs.items())
    assert sha(destinations[0])==sha(destinations[1])
    print(json.dumps({'controles':controls,'salidas':manifest['salidas']},ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
