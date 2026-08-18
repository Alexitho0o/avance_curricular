# -*- coding: utf-8 -*-
# Excel de revision para el usuario - rector_v2 (segunda version)
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = os.path.expanduser("~/mnt/avance_curricular/oferta_academica_2027")
VAL_DIR = os.path.join(BASE, "06_validaciones")
INST = os.path.join(BASE, "03_fuentes_institucionales/entregas_docencia/20260828_etapa2_entrega_02/Oferta 2027 Etapa2_rector_v2.xlsx")
OUT_DIR = os.path.join(BASE, "07_resultados", "reportes_pendientes_etapa2")
os.makedirs(OUT_DIR, exist_ok=True)
STAMP = "20260828"

val2 = json.load(open(os.path.join(VAL_DIR, "VALIDACION_COMPLETA_ETAPA2_20260828_rector_v2.json"), encoding="utf-8"))
concil2 = json.load(open(os.path.join(VAL_DIR, "CONCILIACION_ETAPA2_VS_REPORTE_5912_20260828_rector_v2.json"), encoding="utf-8"))
estructura_cmp = json.load(open(os.path.join(VAL_DIR, "COMPARACION_ESTRUCTURAL_Y_REGISTROS_V1_VS_V2_20260828.json"), encoding="utf-8"))
hallazgos_cmp = json.load(open(os.path.join(VAL_DIR, "COMPARACION_HALLAZGOS_V1_VS_V2_20260828.json"), encoding="utf-8"))
contrato = json.load(open(os.path.join(BASE, "11_gobernanza", "REGLAS_MANUAL_ETAPA2_OFERTA_NUEVA_2027_20260827.json"), encoding="utf-8"))

RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BLUE = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14, color="C00000")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.Workbook()

# ================= Resumen =================
ws1 = wb.active
ws1.title = "Resumen"
ws1["A1"] = "OFERTA ACADEMICA NUEVA 2027 - ETAPA 2 - SEGUNDA VERSION (rector_v2) - 20260828"
ws1["A1"].font = Font(bold=True, size=14)
ws1["A2"] = "Archivo: Oferta 2027 Etapa2_rector_v2.xlsx | Copia gobernada: entregas_docencia/20260828_etapa2_entrega_02/"
bloq_v2 = val2["resumen"].get("BLOQUEANTE", 0)
filas_bloq_v2 = len(set(h["fila_excel"] for h in val2["hallazgos"] if h["severidad"]=="BLOQUEANTE"))
estado = "NO CARGAR"
ws1["A4"] = f"ESTADO: {estado}"
ws1["A4"].font = TITLE_FONT
ws1["A5"] = ("La estructura oficial PES independiente sigue sin confirmarse, persisten contradicciones documentales "
             "no resueltas (VERSION, rango de fechas 2027) y quedan hallazgos BLOQUEANTE. No se genera CSV de carga.")
ws1["A5"].alignment = WRAP
ws1.merge_cells("A5:F5")
ws1.row_dimensions[5].height = 40

r = 7
ws1.cell(row=r, column=1, value="Metrica").font = BOLD
ws1.cell(row=r, column=2, value="Version 1 (20260827)").font = BOLD
ws1.cell(row=r, column=3, value="Version 2 / rector_v2 (20260828)").font = BOLD
r += 1
filas = [
    ("Registros reales", estructura_cmp["estructura"]["v1"]["registros_reales"], estructura_cmp["estructura"]["v2"]["registros_reales"]),
    ("Columnas totales", estructura_cmp["estructura"]["v1"]["columnas_totales"], estructura_cmp["estructura"]["v2"]["columnas_totales"]),
    ("Primeras 48 columnas coinciden con Anexo2 pag.35-41", estructura_cmp["estructura"]["primeras_48_columnas_v1_coinciden_con_anexo2_pag35_41"], estructura_cmp["estructura"]["primeras_48_columnas_v2_coinciden_con_anexo2_pag35_41"]),
    ("Columna SECTOR IPSS presente", estructura_cmp["estructura"]["columna_sector_ipss_en_v1"], estructura_cmp["estructura"]["columna_sector_ipss_en_v2"]),
    ("Hallazgos BLOQUEANTE", hallazgos_cmp["resumen_severidad_v1"].get("BLOQUEANTE",0), hallazgos_cmp["resumen_severidad_v2"].get("BLOQUEANTE",0)),
    ("Hallazgos REVISION_MANUAL", hallazgos_cmp["resumen_severidad_v1"].get("REVISION_MANUAL",0), hallazgos_cmp["resumen_severidad_v2"].get("REVISION_MANUAL",0)),
    ("Hallazgos INFORMATIVO", hallazgos_cmp["resumen_severidad_v1"].get("INFORMATIVO",0), hallazgos_cmp["resumen_severidad_v2"].get("INFORMATIVO",0)),
    ("Filas con al menos un BLOQUEANTE", hallazgos_cmp["filas_con_bloqueante_v1"], hallazgos_cmp["filas_con_bloqueante_v2"]),
    ("Posibles duplicados vs reporte 5912", 18, concil2["resumen"].get("POSIBLE_DUPLICADO_ETAPA1", 0)),
]
for label, a, b in filas:
    ws1.cell(row=r, column=1, value=label)
    ws1.cell(row=r, column=2, value=str(a))
    c3 = ws1.cell(row=r, column=3, value=str(b))
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        if b < a: c3.fill = GREEN
        elif b > a: c3.fill = RED
        else: c3.fill = YELLOW
    r += 1

r += 1
ws1.cell(row=r, column=1, value="Cambios de registros (v1 -> v2):").font = BOLD
r += 1
for k, v in estructura_cmp["resumen_registros"].items():
    ws1.cell(row=r, column=1, value=k)
    ws1.cell(row=r, column=2, value=v)
    r += 1

r += 1
ws1.cell(row=r, column=1, value="Leyenda de colores:").font = BOLD
r += 1
leyenda = [("Rojo","Bloqueo para carga (BLOQUEANTE)", RED), ("Amarillo","Revision manual requerida", YELLOW), ("Verde","Validacion superada / mejora respecto de v1", GREEN), ("Azul","Informacion contextual", BLUE)]
for label, desc, fill in leyenda:
    c = ws1.cell(row=r, column=1, value=label); c.fill = fill
    ws1.cell(row=r, column=2, value=desc)
    r += 1

ws1.column_dimensions["A"].width = 55
ws1.column_dimensions["B"].width = 28
ws1.column_dimensions["C"].width = 32

# ================= Comparacion con version 1 =================
ws2 = wb.create_sheet("Comparación con versión 1")
cols2 = ["campo","v1_bloqueante","v1_revision_manual","v1_informativo","v1_total","v2_bloqueante","v2_revision_manual","v2_informativo","v2_total","estado"]
for j, c in enumerate(cols2, start=1):
    cell = ws2.cell(row=1, column=j, value=c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT
ws2.freeze_panes = "A2"
color_by_estado = {"RESUELTO": GREEN, "PARCIALMENTE_RESUELTO": GREEN, "PERSISTENTE_SIN_CAMBIO": YELLOW, "NUEVO": RED, "EMPEORO": RED, "SIN_HALLAZGOS": None}
i = 2
for row in hallazgos_cmp["comparacion_por_campo"]:
    if row["estado"] == "SIN_HALLAZGOS":
        continue
    for j, c in enumerate(cols2, start=1):
        cell = ws2.cell(row=i, column=j, value=row.get(c))
        cell.border = BORDER
    fill = color_by_estado.get(row["estado"])
    if fill:
        for j in range(1, len(cols2)+1):
            ws2.cell(row=i, column=j).fill = fill
    i += 1
# registros nuevos/eliminados/modificados detalle
i += 1
ws2.cell(row=i, column=1, value="Detalle de registros (comparacion por llave diagnostica)").font = BOLD
i += 1
cols2b = ["fila_v1","fila_v2","sede","carrera","clasificacion","cambios_de_valor"]
for j, c in enumerate(cols2b, start=1):
    cell = ws2.cell(row=i, column=j, value=c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT
i += 1
for reg in estructura_cmp["detalle_registros"]:
    for j, c in enumerate(cols2b, start=1):
        cell = ws2.cell(row=i, column=j, value=reg.get(c))
        cell.border = BORDER
        cell.alignment = WRAP if c == "cambios_de_valor" else Alignment(vertical="top")
    if reg["clasificacion"] == "REGISTRO_NUEVO":
        fill = BLUE
    elif reg["clasificacion"] == "REGISTRO_ELIMINADO":
        fill = RED
    elif reg["clasificacion"] == "REGISTRO_MODIFICADO":
        fill = YELLOW
    else:
        fill = GREEN
    for j in range(1, len(cols2b)+1):
        ws2.cell(row=i, column=j).fill = fill
    i += 1
widths2 = [12,18,16,16,16,16,16,16,16,24]
for j, w in enumerate(widths2, start=1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ================= Detalle de hallazgos =================
ws3 = wb.create_sheet("Detalle de hallazgos")
cols3 = ["fila_excel","llave_diagnostica","carrera","sede","modalidad","jornada","campo","valor_observado","regla_incumplida","fuente","pagina","linea","severidad","accion_requerida"]
for j, c in enumerate(cols3, start=1):
    cell = ws3.cell(row=1, column=j, value=c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT
ws3.freeze_panes = "A2"
for i, f in enumerate(val2["hallazgos"], start=2):
    for j, c in enumerate(cols3, start=1):
        cell = ws3.cell(row=i, column=j, value=f.get(c))
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in ("regla_incumplida","accion_requerida","llave_diagnostica")))
    if f["severidad"] == "BLOQUEANTE":
        fill = RED
    elif f["severidad"] == "REVISION_MANUAL":
        fill = YELLOW
    elif f["severidad"] == "ADVERTENCIA":
        fill = YELLOW
    else:
        fill = BLUE
    for j in range(1, len(cols3)+1):
        ws3.cell(row=i, column=j).fill = fill
widths3 = [10,45,30,22,10,8,26,20,55,32,8,8,14,45]
for j, w in enumerate(widths3, start=1):
    ws3.column_dimensions[get_column_letter(j)].width = w

# ================= Pendientes institucionales =================
ws4 = wb.create_sheet("Pendientes institucionales")
ws4["A1"] = "Contradicciones documentales y pendientes materiales (persisten desde la version 1)"
ws4["A1"].font = Font(bold=True, size=12)
r = 3
for c in contrato["contradicciones_detectadas"]:
    cell = ws4.cell(row=r, column=1, value=c["campo"]); cell.font = BOLD; cell.fill = YELLOW
    r += 1
    ws4.cell(row=r, column=1, value=c["descripcion"]).alignment = WRAP
    ws4.merge_cells(f"A{r}:E{r}"); ws4.row_dimensions[r].height = 45
    r += 1
    ws4.cell(row=r, column=1, value="Resolucion: " + c["resolucion"]).alignment = WRAP
    ws4.merge_cells(f"A{r}:E{r}"); ws4.row_dimensions[r].height = 30
    r += 2

r += 1
ws4.cell(row=r, column=1, value="Pendientes especificos de esta segunda version:").font = BOLD
r += 1
pendientes_v2 = [
    f"18 registros siguen coincidiendo EXACTAMENTE con el reporte 5912 (Etapa 1 validada) -- sin cambio respecto de v1. Responde: Docencia.",
    "COD_SEDE vacio en el registro TECNICO EN FARMACIA / SEDE CONCEPCION (fila 49 en v2): se perdio el codigo de sede que si estaba presente en v1 (fila 5). Responde: Docencia.",
    "Se agrego una columna sin encabezado (posicion 49, antes de SECTOR IPSS) que no existia en v1: la estructura tiene ahora 50 columnas en vez de 49, y SECTOR IPSS se desplazo a la posicion 50. La columna extra sigue sin resolverse. Responde: Docencia / SIES.",
    "VACANTES_SEGUNDO_SEMESTRE=0 con vigencia=1 aumento de 35 a 37 filas. Responde: Docencia.",
    "VERSION paso de estar informado en 40 a 49 filas, bajo una contradiccion del instructivo aun no resuelta. Responde: SIES.",
    "AREA_ADMIN_DERECHO y AREA_TECNO_INFO_COMUNICA mantienen exactamente los mismos 3 hallazgos bloqueantes cada una, sin cambios. Responde: Docencia.",
]
for p in pendientes_v2:
    ws4.cell(row=r, column=1, value="- " + p).alignment = WRAP
    ws4.merge_cells(f"A{r}:E{r}"); ws4.row_dimensions[r].height = 30
    r += 1

ws4.column_dimensions["A"].width = 100

# ================= Reglas aplicadas =================
ws5 = wb.create_sheet("Reglas aplicadas")
cols5 = ["posicion","letra","nombre_oficial","tipo","dominio","obligatorio_o_condicional","pagina","linea","estado_respaldo"]
for j, c in enumerate(cols5, start=1):
    cell = ws5.cell(row=1, column=j, value=c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT
ws5.freeze_panes = "A2"
for i, col in enumerate(contrato["columnas"], start=2):
    for j, c in enumerate(cols5, start=1):
        cell = ws5.cell(row=i, column=j, value=col.get(c))
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c == "dominio"))
    if col["estado_respaldo"] != "RESPALDADO_ESTRUCTURA_MANUAL_PAG35_41":
        for j in range(1, len(cols5)+1):
            ws5.cell(row=i, column=j).fill = YELLOW
    else:
        for j in range(1, len(cols5)+1):
            ws5.cell(row=i, column=j).fill = GREEN
ws5.cell(row=len(contrato["columnas"])+3, column=1, value="Nota: reglas identicas a la Fase 3 de la version 1 (no se recreo el contrato de columnas; ver 11_gobernanza/REGLAS_MANUAL_ETAPA2_OFERTA_NUEVA_2027_20260827.*)").font = Font(italic=True)
widths5 = [8,6,32,10,45,20,8,8,45]
for j, w in enumerate(widths5, start=1):
    ws5.column_dimensions[get_column_letter(j)].width = w

out_path = os.path.join(OUT_DIR, f"REPORTE_REVISION_ETAPA2_RECTOR_V2_{STAMP}.xlsx")
wb.save(out_path)
print("Excel generado:", out_path)
