"""Govern the new institutional source and audit identity before any PES changes."""

import csv
import hashlib
import io
import json
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from preparar_etapa1_octubre_20261020 import FIELDS, ROOT, write_once, tsv


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(text):
    return " ".join("".join(c for c in unicodedata.normalize("NFD", text.upper()) if unicodedata.category(c) != "Mn").split())


def comparison_only(text):
    return re.sub(r"^DIPLOMADO EN ", "DIPLOMADO ", norm(text))


def main():
    xlsx = Path.home() / "Downloads" / "Oferta Diplomados 2027 - SIES vf.xlsx"
    error = Path.home() / "Downloads" / "20-10-2026_18-06-40_Oferta-Académica-Vigente-Editada-TP.csv.txt"
    base = ROOT / "07_resultados/borradores_prueba_no_oficiales/etapa1_octubre_20261020/PRUEBA_PES_OFERTA_2027_ETAPA1_17467_20261020_V04_SIN_SALTO_FINAL.csv"
    inputs = {str(p): sha(p) for p in (xlsx, base)}
    if error.exists():
        inputs[str(error)] = sha(error)
    wb = load_workbook(xlsx, data_only=False)
    assert wb.sheetnames == ["Hoja2"]
    ws = wb["Hoja2"]
    headers = [c.value for c in ws[4]]
    assert headers == ["Período Académico", "Nombre Programa", "N° Vacantes", "$ Matrícula", "$ Titulación", "$ Certificado", "$ Arancel", "Link Sitio Web"]
    doc = []
    for row in ws.iter_rows(min_row=5):
        if not any(c.value is not None for c in row):
            continue
        assert not any(c.data_type == "f" for c in row)
        values = [c.value for c in row]
        assert all(isinstance(v, (int, float)) and not isinstance(v, bool) and v >= 0 for v in values[2:7])
        doc.append({"fila_excel": row[0].row, **dict(zip(headers, values))})
    assert len(doc) == 29 and len({norm(d["Nombre Programa"]) for d in doc}) == 29
    raw = list(csv.reader(io.StringIO(base.read_text(encoding="utf-8")), delimiter=";", strict=True))
    assert len(raw) == 161 and all(len(r) == len(FIELDS) == 48 for r in raw)
    rows = [{"fila_csv": i, **dict(zip(FIELDS, r))} for i, r in enumerate(raw, 1)]
    diplomas = [r for r in rows if r["COD_NIVEL_GLOBAL"] == "3" and r["COD_NIVEL_CARRERA"] == "5"]
    assert len(diplomas) == 21
    if error.exists():
        affected = {int(n) for n in re.findall(r"Errores en linea (\d+):", error.read_text())}
        assert {d["fila_csv"] for d in diplomas} == affected
    classified = []
    matched_doc = set()
    for row in diplomas:
        exact = [d for d in doc if norm(d["Nombre Programa"]) == norm(row["NOMBRE_CARRERA"])]
        candidates = exact or [d for d in doc if comparison_only(d["Nombre Programa"]) == comparison_only(row["NOMBRE_CARRERA"])]
        assert len(candidates) <= 1
        status = "COINCIDENCIA_NOMBRE_SIN_CLAVE_SIES" if exact else ("CANDIDATO_POR_OMISION_DE_EN_NO_APLICADO" if candidates else "SIN_CORRESPONDENCIA_CONFIRMADA")
        d = candidates[0] if candidates else {}
        if d:
            matched_doc.add(d["fila_excel"])
        classified.append([row["fila_csv"], row["COD_SEDE"], row["COD_CARRERA"], row["MODALIDAD"], row["COD_JORNADA"], row["VERSION"], row["NOMBRE_CARRERA"], row["VIGENCIA_CARRERA"], status, d.get("fila_excel", ""), d.get("Nombre Programa", ""), d.get("Período Académico", ""), d.get("N° Vacantes", ""), d.get("$ Matrícula", ""), d.get("$ Titulación", ""), d.get("$ Certificado", ""), d.get("$ Arancel", ""), "Confirmar vigencia/identidad y vacantes por semestre; no se aplican cambios"])
    counts = Counter(r[8] for r in classified)
    assert sorted(counts.values()) == [2, 3, 16]
    delivery = ROOT / "03_fuentes_institucionales/entregas_docencia/20261023_diplomados_vf"
    audit = ROOT / "06_validaciones/etapa1_octubre_20261023_diplomados"
    gov = ROOT / "11_gobernanza/etapa1_octubre_20261023_diplomados"
    write_once(delivery / xlsx.name, xlsx.read_bytes())
    if error.exists():
        write_once(gov / error.name, error.read_bytes())
    tsv(audit / "CRUCE_21_DIPLOMADOS.tsv", [["fila_csv", "cod_sede", "cod_carrera", "modalidad", "jornada", "version", "nombre_sies", "vigencia_sies", "estado_cruce", "fila_excel", "nombre_docencia", "periodo", "vacantes_sin_desglose_semestral", "matricula", "titulacion", "certificado", "arancel", "pendiente"]] + classified)
    tsv(audit / "PROGRAMAS_EXCEL_SIN_CORRESPONDENCIA.tsv", [["fila_excel"] + headers] + [[d["fila_excel"]] + [d[h] for h in headers] for d in doc if d["fila_excel"] not in matched_doc])
    note = """# Revision de entrega de diplomados: carga pendiente

23-10-2026. Oferta Academica 2027, segundo proceso, Etapa 1, ID 17467.
Base: V04, 161 registros, 48 campos, fechas corregidas y sin salto final.
Entrega: Oferta Diplomados 2027 - SIES vf.xlsx, Hoja2, encabezado fila 4.

El Excel contiene 29 programas, 15 vacantes por programa y no contiene codigos
SIES, sede, modalidad, jornada, version ni vigencia. El rechazo PES afecta a
21 diplomados. Hay 3 coincidencias de nombre tras comparar sin tildes/caja,
2 candidatos por la presencia/ausencia de EN y 16 sin correspondencia por
estos criterios. Estas comparaciones no renombraron datos ni autorizaron joins.
Los 24 programas restantes del Excel no se declaran nuevos automaticamente.

Bloqueos para generar carga:
1. Confirmar si la nueva entrega sustituye o complementa la oferta y aportar
   identidad SIES o correspondencia autorizada. La ausencia en Excel no prueba
   que un programa deba cambiar de vigencia; se requieren estados confirmados.
2. Definir vacantes por semestre. Periodos 2027-00 (12), 2027-01 (8),
   2027-02 (7), 2027-03 (2); no existe equivalencia semestral en el libro.

Importes observados: matricula 400000 (10 programas), 600000 (2), 990000 (17).
Titulacion, certificado y arancel son cero en los 29 registros. No intercambiar
matricula con arancel. Total institucional informado: 435 cupos; no sumarlos
a los 5.461 de la base sin resolver identidad y cobertura de etapas.

El archivo de error del 20-10-2026 18:06:40 no se encuentra actualmente en
Descargas, Escritorio ni el subproyecto. Sus mensajes se conocen por la lectura
registrada en esta conversacion; no se reconstruye ni presenta como original.

El rechazo PES establece 0 en tres cobros si no corresponden, y rechaza
vigencia 1 con ambas vacantes cero. No se inventan vacantes ni estados para
sortear la validacion. No se genero nuevo CSV de carga ni se modifico V04.

Reanudacion: recibir correspondencias/vigencias y distribucion semestral,
aplicar unicamente datos autorizados a la base, generar diff por celda,
validar estructura y reglas, conservar UTF-8, punto y coma y ausencia de
salto final; guardar copia en Escritorio. Aceptacion final pendiente de PES.
Respaldo local reproducible. No commit, push ni carga a PES.
"""
    write_once(gov / "BLOQUEOS_Y_REANUDACION.md", note.encode("utf-8"))
    manifest = {
        "fecha": datetime.now(ZoneInfo("America/Santiago")).isoformat(),
        "estado": "BLOQUEADO_IDENTIDAD_VIGENCIA_Y_DISTRIBUCION_SEMESTRAL",
        "entradas": [{"ruta": p, "sha256": h} for p, h in inputs.items()],
        "script": {"ruta": str(Path(__file__).resolve()), "sha256": sha(Path(__file__))},
        "controles": {"registros_base": 161, "diplomados_pes": 21, "programas_docencia": 29, "cruce": dict(counts), "periodos": dict(Counter(d["Período Académico"] for d in doc)), "vacantes_excel": sum(d["N° Vacantes"] for d in doc)},
        "evidencia_rechazo": "Archivo original disponible" if error.exists() else "Lectura previa en conversacion; original ya no disponible en las rutas revisadas",
        "archivos": [{"ruta": str(p), "sha256": sha(p)} for p in [delivery / xlsx.name, audit / "CRUCE_21_DIPLOMADOS.tsv", audit / "PROGRAMAS_EXCEL_SIN_CORRESPONDENCIA.tsv", gov / "BLOQUEOS_Y_REANUDACION.md"]],
        "salida_carga": None,
    }
    if not (gov / "MANIFIESTO.json").exists():
        write_once(gov / "MANIFIESTO.json", (json.dumps(manifest, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
    assert all(sha(Path(p)) == h for p, h in inputs.items())
    assert sha(delivery / xlsx.name) == inputs[str(xlsx)]
    print(json.dumps(manifest["controles"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
