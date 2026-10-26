"""Read-only reconciliation against accepted V06; prepare an assessment, not a load."""

import csv
import io
import json
import re
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from preparar_etapa1_octubre_20261020 import ROOT, FIELDS, sha, tsv, write_once


def norm(value):
    text = ''.join(c for c in unicodedata.normalize('NFD', value.upper()) if not unicodedata.combining(c))
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', text).split())


def main():
    frozen = ROOT / '07_resultados/cargas_congeladas/20261023_etapa1_octubre_v06_aceptada'
    baseline = frozen / 'CARGA_OFERTA_2027_ETAPA1_17467_20261023_V06.csv'
    decisions = frozen / 'DECISIONES_21_DIPLOMADOS.tsv'
    doc = ROOT / '03_fuentes_institucionales/entregas_docencia/20261023_diplomados_vf/Oferta Diplomados 2027 - SIES vf.xlsx'
    manual = ROOT / '00_fuentes_originales/etapa1_octubre_20261020/Instructivo_Oferta_Academica_2027_CFT_E_IP_DFE.pdf'
    inputs = {str(p): {'sha256': sha(p), 'modo': oct(p.stat().st_mode & 0o777)} for p in sorted(frozen.iterdir()) if p.is_file()}
    inputs.update({str(p): {'sha256': sha(p)} for p in [doc, manual]})
    assert sha(baseline) == 'd7fe88035cd1215d448272a9854bee0932c45c200f4f80bb58636b62a6ae4e0b'
    rows = list(csv.reader(io.StringIO(baseline.read_text()), delimiter=';'))
    assert len(rows) == 161 and all(len(r) == 48 for r in rows)
    sies = [dict(zip(FIELDS, r)) for r in rows]
    diplomas = {r['COD_CARRERA']: r for r in sies if r['COD_NIVEL_GLOBAL'] == '3' and r['COD_NIVEL_CARRERA'] == '5'}
    assert len(diplomas) == 21
    with decisions.open(newline='') as f:
        accepted = {int(r['fila_excel']): r for r in csv.DictReader(f, delimiter='\t') if r['vigencia'] == '1'}
    assert len(accepted) == 7
    pages = PdfReader(manual).pages
    schema = []
    for p in range(35, 42):
        schema.extend(re.findall(r'^\s*(?:[A-Z]|A[A-V])\s+([A-Z][A-Z_]+)\b', pages[p-1].extract_text(), re.MULTILINE))
    assert schema == FIELDS, 'El contrato debe seguir el Anexo 2 actual'

    ws = load_workbook(doc, data_only=False)['Hoja2']
    assert (ws.max_row, ws.max_column) == (33, 8)
    programs = []
    name_notes = {
        8: 'Programa distinto de Gobernanza de Datos (codigo 96). No fusionar por palabras compartidas.',
        13: 'Revisar denominacion: la fuente escribe "Analytics y business inteligence". No se corrige automaticamente.',
        20: 'Confirmar denominacion sin abreviaciones ni barra UX/UI, conforme a p.35.',
        27: 'Confirmar denominacion completa de TI, sin expandir siglas por inferencia (p.35).',
        29: 'Confirmar nombre: la fuente repite "Diplomado en Diplomado en" y contiene punto final.',
        31: 'Adecuar punto final y espacio no separable al formato de nombre de p.35; conservar original en fuente.',
    }
    for n in range(5, 34):
        raw = [ws.cell(n, c).value for c in range(1, 9)]
        assert all(ws.cell(n, c).data_type != 'f' for c in range(1, 9))
        period, name, vac, amount, tit, cert, arancel, url = raw
        assert vac == 15 and tit == cert == arancel == 0
        a = accepted.get(n)
        if a:
            assert a['nombre_docencia'] == name
            r = diplomas[a['cod_carrera']]
            assert r['VIGENCIA_CARRERA'] == '1' and r['NOMBRE_CARRERA'] == a['nombre_sies']
            assert r['ARANCEL_ANUAL'] == str(amount) and r['VALOR_MATRICULA_ANUAL'] == '0'
            state, note = 'YA EN ETAPA 1', a['respaldo']
        else:
            state = 'CANDIDATO ETAPA 2'
            assert not any(norm(name) == norm(r['NOMBRE_CARRERA']) for r in sies if r['VIGENCIA_CARRERA'] in ['1','2'])
            note = name_notes.get(n, 'Completar antecedentes academicos y confirmar oferta 2027. No copiar datos de nombres similares.')
        programs.append([name, state, a['cod_carrera'] if a else '', a['nombre_sies'] if a else '', note, n, period, vac, amount, tit, cert, arancel, url or ''])
    candidates = [p for p in programs if p[1] == 'CANDIDATO ETAPA 2']
    assert len(candidates) == 22 and len({norm(p[0]) for p in programs}) == 29

    # The field review is specific to postitulo in the October manual, not August rules.
    missing = {
        'COD_SEDE': ('35, 32', 'Confirmar sede de imparticion; codigo numerico del catalogo SIES.'),
        'NOMBRE_SEDE': ('35, 32', 'Usar nombre exacto de sede SIES. Mayusculas A-Z sin acentos, numeros, sin abreviaciones.'),
        'MODALIDAD': ('35, 56', '1 presencial; 2 semipresencial; 3 no presencial. Parejas modalidad/jornada: 1/1, 1/2, 2/3, 3/4.'),
        'COD_JORNADA': ('35, 56', '1 diurna; 2 vespertina; 3 semipresencial; 4 a distancia. Validar pareja con modalidad.'),
        'VERSION': ('36', 'Obligatoria, numero del 1 al 9 segun versiones existentes. No asignar 1 por defecto.'),
        'COD_TIPO_PLAN_CARRERA': ('36', '1 regular; 2 especial; 3 regular de continuidad.'),
        'DURACION_ESTUDIOS': ('36, 57', 'Duracion de estudios en semestres. <=14; si supera, gestionar con encargado SIES. No deducir de periodo Excel.'),
        'DURACION_TITULACION': ('36, 57', 'Duracion de titulacion en semestres. No confundir costo cero con duracion cero.'),
        'DURACION_TOTAL': ('36, 57', 'Semestres totales. Estudios <= total <= estudios + titulacion.'),
        'REGIMEN': ('36', '1 semestres; 2 trimestres; 3 anos; 4 bimestres.'),
        'DURACION_REGIMEN': ('36-37', 'Duracion formal segun malla en unidades del regimen declarado.'),
        'NOMBRE_TITULO': ('37', 'Nombre de certificacion final. Mayusculas A-Z sin acentos, numeros, sin abreviaciones. No copiar nombre de programa sin respaldo.'),
        'ANIO_INICIO': ('37', 'Ano en que el programa inicio actividades. Periodo de oferta 2027 no acredita ano de inicio.'),
        'ACREDITACION': ('37', 'Del programa, no institucional: 1 si, 2 no, 3 en proceso. Confirmar; la mencion de precarga 2 no acredita el dato del programa nuevo.'),
        'REQUISITO_INGRESO': ('37', 'Nivel minimo: 1 media; 2 tecnico superior; 3 bachillerato; 4 ciclo basico; 5 plan comun; 6 titulo profesional.'),
        'SEMESTRES_RECONOCIDOS': ('37', 'Cantidad reconocida como condicion de acceso. No imputar cero sin confirmacion.'),
        'AREA_ACTUAL': ('38', 'Codigo CINE F-2013 de 1 a 10. Clasificacion institucional; no inferir por palabras del nombre.'),
    }
    rules = []
    for i, field in enumerate(FIELDS):
        letter = get_column_letter(i+1)
        if field in missing:
            page, rule = missing[field]
            state, count = 'FALTA DATO', 22
        elif field == 'COD_CARRERA':
            page, rule, state, count = '31, 35', 'Dejar vacio. SIES asignara codigo despues del cierre. No reutilizar codigos dados de baja.', 'VACIO EXIGIDO', 0
        elif field == 'NOMBRE_CARRERA':
            page, rule, state, count = '35', 'Nombre disponible. Preparar mayusculas A-Z sin acentos y numeros, sin abreviaciones. Revisar denominaciones marcadas; no renombrar originales.', 'REVISAR FORMATO Y NOMBRE', 22
        elif field == 'CARACTERISTICAS_TIPO_PLAN':
            page, rule, state, count = '36', 'Si tipo 1 o 3, NO APLICA; si tipo 2, descripcion de caracteristica especial. Depende del tipo aun no informado.', 'DEPENDE DEL TIPO DE PLAN', 22
        elif field in ['COD_NIVEL_GLOBAL', 'COD_NIVEL_CARRERA']:
            page, rule, state, count = '37, 57', 'Diplomado informable: global 3 y especifico 5, desde un semestre. Confirmar duracion y clasificacion antes de asignar.', 'CONFIRMAR CLASIFICACION', 22
        elif 22 <= i <= 31:
            page, rule, state, count = '9-10, 38-39', 'Para postitulos, cero en todas las areas de destino. Aplicar una vez confirmada clasificacion; no son faltantes a pedir a Docencia.', 'CERO SEGUN NIVEL', 0
        elif field.startswith('VACANTES_'):
            page, rule, state, count = '10, 40', 'Numero por semestre; suma anual >0 para vigencia 1. Excel informa una sola cantidad 15; criterio usuario previo 15+15. Ratificar alcance a los 22 candidatos.', 'CONFIRMAR ALCANCE DEL CRITERIO', 22
        elif 34 <= i <= 41:
            page, rule, state, count = '40-41', 'Columna de obligatoriedad aplicable a Postitulo: Obligatorio / Vacio. No trasladar enlaces ni datos de admision a estos campos.', 'VACIO EXIGIDO', 0
        elif field == 'FORMATO_VALOR':
            page, rule, state, count = '41', '1 pesos chilenos; 2 UF. Criterio CLP de V06 propuesto para nuevos; ratificar alcance.', 'CONFIRMAR ALCANCE DEL CRITERIO', 22
        elif field in ['VALOR_MATRICULA_ANUAL', 'ARANCEL_ANUAL']:
            page, rule, state, count = '41', 'Importe obligatorio, numerico sin separadores. Usuario corrigio V06: columna D de Docencia es arancel y matricula 0. Conservar original y ratificar alcance a nuevos.', 'CONFIRMAR ALCANCE DEL CRITERIO', 22
        elif field in ['COSTO_TITULACION', 'VALOR_CERTIFICADO_DIPLOMA']:
            page, rule, state, count = '41', 'Importe obligatorio numerico sin separadores. Docencia informa 0 en todos los candidatos, no es un vacio.', 'DATO DISPONIBLE', 0
        elif field == 'VIGENCIA_CARRERA':
            page, rule, state, count = '41', '1 vigente con nuevos; 2 vigente sin nuevos; 3 no vigente. Confirmar oferta efectiva y universo de nuevos. Candidato no equivale a autorizacion final.', 'CONFIRMAR OFERTA', 22
        else:
            raise AssertionError(field)
        rules.append([letter, field, state, count, rule, page])
    assert len(rules) == 48 and len(missing) == 17

    tag = 'preparacion_etapa2_noviembre_20261023'
    out = ROOT / '10_entrega' / tag
    gov = ROOT / '11_gobernanza' / tag
    audit = ROOT / '06_validaciones' / tag
    headers = ['Programa Docencia', 'Tratamiento', 'Codigo SIES V06', 'Programa SIES V06', 'Observacion', 'Fila Hoja2', 'Periodo original', 'Vacantes originales', 'Col. D: Matricula original', 'Titulacion original', 'Certificado original', 'Arancel original', 'Web original']
    rule_headers = ['Columna', 'Campo SIES', 'Situacion', 'Candidatos pendientes', 'Regla y accion requerida', 'Paginas manual']
    tsv(audit / 'CONCILIACION_29_PROGRAMAS.tsv', [headers] + programs)
    tsv(gov / 'REVISION_48_COLUMNAS.tsv', [rule_headers] + rules)
    tsv(audit / 'FALTANTES_POR_PROGRAMA.tsv', [['fila_docencia', 'programa', 'campo', 'pagina', 'motivo']] + [[p[5], p[0], field, page, why] for p in candidates for field, (page, why) in missing.items()])
    tsv(audit / 'BAJAS_ETAPA1_SIN_REACTIVACION.tsv', [['codigo', 'nombre', 'vigencia']] + [[c, r['NOMBRE_CARRERA'], r['VIGENCIA_CARRERA']] for c, r in diplomas.items() if r['VIGENCIA_CARRERA'] == '3'])
    report = '''# Preparacion Etapa 2 de noviembre: diplomados 2027

Proceso 2026, referencia 2027. Analisis al 23-10-2026. NO ES ARCHIVO DE CARGA.

## Resultado
29 programas de Docencia, 7 representados en Etapa 1 V06 y 22 candidatos nuevos o faltantes.
Los 7 se excluyen de la propuesta: codigos 53, 66, 68, 69, 70, 96 y 99.
Se conservan las equivalencias ya decididas para 68 y 99, aunque los nombres difieren.
Los 14 codigos con vigencia 3 permanecen sin cambios. No se convierten automaticamente
en 14 altas nuevas ni se reutilizan sus codigos. El universo se construye desde Docencia.
22 es cantidad de programas fuente, no numero definitivo de filas PES: sede, modalidad,
jornada y versiones todavia no estan informadas y pueden cambiar el detalle de registros.

## Pendientes academicos por cada candidato
17 campos sin dato directo: sede y nombre, modalidad, jornada, version, tipo de plan,
duracion de estudios/titulacion/total, regimen y duracion del regimen, certificacion final,
ano real de inicio, acreditacion del programa, requisito de ingreso, semestres reconocidos
y area actual. Son 374 ausencias de datos, no 374 rechazos observados en PES.
Caracteristicas del plan depende del tipo: NO APLICA para 1/3; describir si 2.
Confirmar duracion desde un semestre para clasificar como diplomado (global 3, nivel 5).
Confirmar nombres observados: filas 13, 20, 27, 29 y 31. No se corrigieron fuentes.
Gestion Estrategica y Gobernanza de Datos (fila 8) no se fusiona con Gobernanza de Datos (96).

## Datos disponibles y decisiones previas
El Excel contiene una sola cifra de 15 vacantes por programa, no un reparto semestral.
La instruccion anterior 15+15 produjo 30 vacantes anuales por diploma V06. Ratificar alcance
a los 22 candidatos; no inferir semestres desde periodos 2027-00/01/02/03.
La fuente conserva importes en Matricula y Arancel=0. Usuario confirmo para V06 que esos
importes son ARANCEL y la matricula queda en 0. Ratificar alcance a nuevos, sin modificar fuente.
Titulacion y certificado son 0 informados, no faltantes. Moneda CLP debe conservar trazabilidad.
No se genera total de vacantes de carga ni importes finales mientras estos alcances esten pendientes.

## Reglas que evitan pedir informacion innecesaria
COD_CARRERA vacio, sin codigo institucional ni encabezados en el futuro CSV (p.31).
VERSION obligatoria del 1 al 9 (p.36): no confundir con el codigo de carrera.
Diez areas de destino en 0 para postitulos (pp.9-10).
AI:AP vacios para postitulos (pp.40-41), incluso ENLACE_INFO_PROGRAMA: los enlaces
ausentes en Docencia no bloquean esta carga. Los existentes se conservan como referencia.
No trasladar las exigencias de pregrado ni las reglas de la Etapa 2 de agosto.

## Preparacion y cierre
Ventana: 02 a 13 de noviembre de 2026 (p.29). Carga ID 17469 (p.30).
48 columnas A:AV verificadas contra Anexo 2, pp.35-41. Falta incorporar estructura vigente
descargada de PES para verificar serializacion final; p.30 dice CSV delimitado por comas.
No asumir automaticamente el delimitador de Etapa 1 ni usar estructura de agosto.
Falta reporte 6048 para conciliacion contra oferta efectivamente validada en PES;
la base actual es V06 aceptada segun confirmacion del usuario, no consulta directa a PES.
No se localizaron archivos con 6048 o 17469 en los nombres del subproyecto ni Descargas
(busqueda dirigida; no prueba de inexistencia bajo otros nombres).
Al completar datos, validar combinaciones modalidad/jornada (p.56), niveles y duraciones
(p.57), nombres y estructura. Confirmar Finalizado y luego conciliar reporte 6049 (p.33).
No se generaron CSV, no se cambio vigencia ni se reabrieron archivos congelados para escritura.

## Reanudacion
Recibir antecedentes academicos de los 22 candidatos, resolver denominaciones y confirmar
alcance de criterios de vacantes e importes; incorporar 6048 y estructura 17469; repetir
conciliacion y validacion antes de generar la carga. No copiar valores de programas parecidos.
'''
    write_once(out / 'RESUMEN_Y_PENDIENTES.md', report.encode())
    payload = {'programas': programs, 'encabezados': headers, 'reglas': rules, 'encabezados_reglas': rule_headers,
               'fuentes': inputs, 'manual': str(manual), 'out': str(out), 'gov': str(gov),
               'controles': {'programas_docencia':29, 'ya_etapa1':7, 'candidatos':22, 'bajas_etapa1_intactas':14,
                            'campos_sin_dato_directo':17, 'ausencias_programa_campo':374,
                            'importes_columna_D_candidatos':dict(Counter(p[8] for p in candidates)),
                            'duplicados_nominales_docencia':0, 'version_congelada_sin_cambios':True}}
    write_once(audit / 'ANALISIS.json', (json.dumps(payload, ensure_ascii=False, indent=2)+'\n').encode())
    for p, info in inputs.items():
        assert sha(Path(p)) == info['sha256']
        if 'modo' in info:
            assert oct(Path(p).stat().st_mode & 0o777) == info['modo']
    print(json.dumps(payload['controles'], ensure_ascii=False))


if __name__ == '__main__':
    main()
