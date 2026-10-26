"""Prepare two decision views; do not modify the PES file."""
import csv
import io
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '04_scripts'))
from conciliar_nombres_diplomados_20261023 import core, sha

source = ROOT / '00_fuentes_originales/etapa1_octubre_20261020/reporte_dinamico_6047_2026-10-20_09:05:04.csv'
docfile = ROOT / '03_fuentes_institucionales/entregas_docencia/20261023_diplomados_vf/Oferta Diplomados 2027 - SIES vf.xlsx'
base = ROOT / '07_resultados/borradores_prueba_no_oficiales/etapa1_octubre_20261020/PRUEBA_PES_OFERTA_2027_ETAPA1_17467_20261020_V04_SIN_SALTO_FINAL.csv'
manual = source.parent / 'Instructivo_Oferta_Academica_2027_CFT_E_IP_DFE.pdf'
sies = [r for r in csv.DictReader(io.StringIO(source.read_bytes().decode('cp1252')), delimiter=';') if r['COD_NIVEL_GLOBAL']=='3']
ws = load_workbook(docfile, data_only=False)['Hoja2']
doc = [{'fila':i, 'nombre':ws.cell(i,2).value, 'matricula':ws.cell(i,4).value} for i in range(5,ws.max_row+1) if ws.cell(i,2).value]
a=[]
for s in sies:
    matches=[d for d in doc if core(d['nombre'])==core(s['NOMBRE_CARRERA'])]
    candidates=matches or [d for d in doc if set(core(s['NOMBRE_CARRERA']).split()) < set(core(d['nombre']).split())]
    state='COINCIDE' if matches else 'POR CONFIRMAR' if candidates else 'SIN CORRESPONDENCIA'
    right=candidates[0]['nombre'] if len(candidates)==1 else f"{len(candidates)} variantes de Inteligencia Artificial; ver Opcion B" if candidates else 'Sin nombre correspondiente en el Excel'
    note='Mismo nombre normalizado; conservar codigo y nombre SIES.' if matches else 'Coincidencia parcial; no autoriza trasladar datos.' if candidates else 'No cambiar vigencia por ausencia en el Excel.'
    if s['COD_CARRERA']=='63':note='Arquitectura no es Infraestructura. Infraestructura Cloud coincide con codigo 69.'
    if s['COD_CARRERA']=='40':note='Docencia Clinica para Tecnicos y Profesionales de la Salud es otro nombre; no hay equivalencia confirmada.'
    if s['COD_CARRERA']=='96':note='Coincide con Gobernanza de Datos ($990.000), no con Gestion Estrategica y Gobernanza de Datos ($400.000).'
    a.append([s['COD_CARRERA'],s['NOMBRE_CARRERA'],right,state,candidates[0]['matricula'] if len(candidates)==1 else None,None,note,', '.join(str(d['fila']) for d in candidates) or None])
a.sort(key=lambda r:({'COINCIDE':0,'POR CONFIRMAR':1,'SIN CORRESPONDENCIA':2}[r[3]],int(r[0])))
b=[]
for d in doc:
    exact=[s for s in sies if core(s['NOMBRE_CARRERA'])==core(d['nombre'])]
    partial=[s for s in sies if set(core(s['NOMBRE_CARRERA']).split()) < set(core(d['nombre']).split())] if not exact else []
    matches=exact or partial
    state='YA IDENTIFICADO EN ETAPA 1' if exact else 'CORRESPONDENCIA POR CONFIRMAR' if partial else 'SIN CORRESPONDENCIA NOMINAL'
    note='Actualizar registro existente; no agregar otro.' if exact else 'Confirmar identidad. Si es un programa distinto, preparar en Etapa 2; no incorporarlo como nuevo en Etapa 1.'
    if d['fila']==8:note='SIES 96 ya coincide con Gobernanza de Datos, fila 32, por $990.000. No sustituirlo por esta variante sin confirmacion.'
    b.append([d['fila'],d['nombre'],' / '.join(s['NOMBRE_CARRERA'] for s in matches) or 'Sin correspondencia nominal',', '.join(s['COD_CARRERA'] for s in matches) or None,d['matricula'],state,None,note])
b.sort(key=lambda r:(r[4],r[0]))
assert len(a)==21 and len(b)==29 and sum(r[3]=='COINCIDE' for r in a)==5
payload={'a':a,'b':b,'inputs':[{ 'ruta':str(p),'sha256':sha(p)} for p in (source,docfile,base,manual)]}
Path(__file__).with_name('opciones.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
print('Opcion A: 21 registros. Opcion B: 29 programas. CSV intacto.')
