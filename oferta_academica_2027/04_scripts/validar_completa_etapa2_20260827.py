# -*- coding: utf-8 -*-
# Fase 5: Validacion completa fila por fila - Etapa 2 (Oferta Academica Nueva TP Adscritas)
# No corrige, no imputa, no elimina registros. Solo detecta y clasifica hallazgos.
import json, os, re, unicodedata
import openpyxl

BASE = os.path.expanduser("~/mnt/avance_curricular/oferta_academica_2027")
INST = os.path.join(BASE, "03_fuentes_institucionales/entregas_docencia/20260827_etapa2_entrega_01/Oferta 2027 Etapa2_prueba_1.xlsx")
OUT = os.path.join(BASE, "06_validaciones")
STAMP = "20260827"
MANUAL = "Instructivo_Oferta_Academica-Acceso_CFT-IP 2027.txt"

AREA_MAP = {
    "1": "AREA_ADMIN_DERECHO", "2": "AREA_AGRI_SILVI_PESCA_VET", "3": "AREA_ARTES_HUMANIDADES",
    "4": "AREA_CIENCIAS_NAT_MAT_ESTAD", "5": "AREA_CS_SOCIAL_PERIODISMO_INFO", "6": "AREA_EDUCACION",
    "7": "AREA_INGE_INDUSTRIA_CONSTRUC", "8": "AREA_SALUD_BIENESTAR", "9": "AREA_SERVICIOS",
    "10": "AREA_TECNO_INFO_COMUNICA",
}
AREA_COLS = list(AREA_MAP.values())

def norm(s):
    if s is None: return ""
    s = str(s).strip().upper()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s)

def is_blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")

def num(v):
    try:
        return float(v)
    except Exception:
        return None

wb = openpyxl.load_workbook(INST, data_only=True)
wb_fmt = openpyxl.load_workbook(INST, data_only=False)
ws = wb["Hoja4"]; ws_fmt = wb_fmt["Hoja4"]
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

rows = []
for r in range(2, 52):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if all(v is None for v in vals):
        continue
    row = dict(zip(headers, vals))
    row["_fila_excel"] = r
    rows.append(row)

# celdas rojas por fila (marca institucional propia)
red_by_row = {}
for r in range(2, 52):
    cols_rojas = []
    for c in range(1, ws.max_column + 1):
        cell = ws_fmt.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.patternType:
            try:
                rgb = fill.fgColor.rgb
            except Exception:
                rgb = None
            if isinstance(rgb, str) and rgb == "FFFF0000":
                cols_rojas.append(headers[c-1])
    if cols_rojas:
        red_by_row[r] = cols_rojas

# marca de color de tema (no RGB) por fila
theme_rows = set()
for r in range(2, 52):
    n = 0
    for c in range(1, ws.max_column + 1):
        cell = ws_fmt.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.patternType:
            try:
                rgb = fill.fgColor.rgb
                is_str = isinstance(rgb, str)
            except Exception:
                is_str = False
            if not is_str:
                n += 1
    if n >= 10:
        theme_rows.add(r)

findings = []

def add(fila, row, campo, valor, regla, severidad, fuente, pagina, linea, accion):
    llave = f"{row.get('NOMBRE_SEDE')} | {row.get('NOMBRE_CARRERA')} | MOD={row.get('MODALIDAD')} | JOR={row.get('COD_JORNADA')} | TIPO_PLAN={row.get('COD_TIPO_PLAN_CARRERA')} | NIVEL={row.get('COD_NIVEL_CARRERA')} | TIT={row.get('NOMBRE_TITULO')}"
    findings.append({
        "fila_excel": fila, "llave_diagnostica": llave,
        "carrera": row.get("NOMBRE_CARRERA"), "sede": row.get("NOMBRE_SEDE"),
        "modalidad": row.get("MODALIDAD"), "jornada": row.get("COD_JORNADA"),
        "campo": campo, "valor_observado": valor, "regla_incumplida": regla,
        "fuente": fuente, "pagina": pagina, "linea": linea,
        "severidad": severidad, "accion_requerida": accion,
    })

# ---- deteccion de duplicados internos por llave diagnostica ----
key_seen = {}
for row in rows:
    k = (norm(row.get("NOMBRE_SEDE")), norm(row.get("NOMBRE_CARRERA")), str(row.get("MODALIDAD")), str(row.get("COD_JORNADA")), str(row.get("COD_TIPO_PLAN_CARRERA")), str(row.get("COD_NIVEL_CARRERA")), norm(row.get("NOMBRE_TITULO")))
    key_seen.setdefault(k, []).append(row["_fila_excel"])

for row in rows:
    fila = row["_fila_excel"]
    modalidad = str(row.get("MODALIDAD")) if row.get("MODALIDAD") is not None else None
    jornada = str(row.get("COD_JORNADA")) if row.get("COD_JORNADA") is not None else None
    tipoplan = str(row.get("COD_TIPO_PLAN_CARRERA")) if row.get("COD_TIPO_PLAN_CARRERA") is not None else None
    nivel = str(row.get("COD_NIVEL_CARRERA")) if row.get("COD_NIVEL_CARRERA") is not None else None
    vigencia = str(row.get("VIGENCIA_CARRERA")) if row.get("VIGENCIA_CARRERA") is not None else None
    regimen = str(row.get("REGIMEN")) if row.get("REGIMEN") is not None else None

    # 1. VIGENCIA_CARRERA vacio
    if is_blank(row.get("VIGENCIA_CARRERA")):
        add(fila, row, "VIGENCIA_CARRERA", row.get("VIGENCIA_CARRERA"), "Campo obligatorio (Anexo2 pag.41) no puede quedar vacio.", "BLOQUEANTE", MANUAL, 41, 1489, "Completar VIGENCIA_CARRERA con 1 o 3 antes de continuar.")
    elif vigencia not in ("1", "3"):
        add(fila, row, "VIGENCIA_CARRERA", row.get("VIGENCIA_CARRERA"), "Para instituciones adscritas en Etapa2 solo se aceptan 1 o 3. Si aparece 2, el Anexo4 (pag.64) indica consultar a SIES.", "REVISION_MANUAL", MANUAL, 64, 2262, "Registrar consulta formal a SIES antes de cargar este registro.")

    # 2. COD_CARRERA no vacio (no deberia ocurrir, pero se valida)
    if not is_blank(row.get("COD_CARRERA")):
        add(fila, row, "COD_CARRERA", row.get("COD_CARRERA"), "COD_CARRERA debe quedar en blanco en Etapa 2 (Anexo2 pag.35; Anexo4 pag.62: 'no debe ser cargado en esta etapa').", "BLOQUEANTE", MANUAL, 35, 1180, "Vaciar el campo antes de cualquier carga.")

    # 3. VERSION -- contradiccion no resuelta
    if not is_blank(row.get("VERSION")):
        add(fila, row, "VERSION", row.get("VERSION"), "Contradiccion no resuelta del propio instructivo: pag.35 (linea 1203) exige VERSION Obligatorio/Modificable; Anexo4 pag.59/64 declara 'La VERSION no corresponde.'", "REVISION_MANUAL", MANUAL, 35, 1203, "Consultar a SIES el tratamiento de VERSION antes de decidir si se carga o se vacia.")

    # 4. RECONOCIMIENTOS_APREN_PREVIOS vacio cuando vigencia=1
    if vigencia == "1" and is_blank(row.get("RECONOCIMIENTOS_APREN_PREVIOS")):
        add(fila, row, "RECONOCIMIENTOS_APREN_PREVIOS", row.get("RECONOCIMIENTOS_APREN_PREVIOS"), "Anexo4 pag.61: 'Cuando la VIGENCIA_CARRERA es 1, se debe completar el RECONOCIMIENTOS_APREN_PREVIOS'.", "BLOQUEANTE", MANUAL, 61, None, "Completar con SI o NO segun corresponda antes de la carga.")

    # 5. Modalidad/Jornada compatibilidad
    compat = {"1": {"1", "2"}, "2": {"3"}, "3": {"4"}}
    if modalidad in compat and jornada is not None and jornada not in compat[modalidad]:
        add(fila, row, "COD_JORNADA", jornada, f"MODALIDAD={modalidad} exige COD_JORNADA en {sorted(compat[modalidad])} (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Corregir jornada o modalidad para que sean compatibles, o confirmar con la institucion.")

    # 6. COD_NIVEL_GLOBAL debe ser 1
    if str(row.get("COD_NIVEL_GLOBAL")) != "1":
        add(fila, row, "COD_NIVEL_GLOBAL", row.get("COD_NIVEL_GLOBAL"), "Para Oferta-Acceso adscritos solo corresponde Pregrado (valor 1).", "BLOQUEANTE", MANUAL, 37, 1260, "Confirmar nivel del programa; no corresponde a Etapa 2 si no es pregrado.")

    # 7. COD_TIPO_PLAN_CARRERA=1 => SEMESTRES_RECONOCIDOS=0
    sem_rec = row.get("SEMESTRES_RECONOCIDOS")
    if tipoplan == "1" and num(sem_rec) not in (0, 0.0):
        add(fila, row, "SEMESTRES_RECONOCIDOS", sem_rec, "Anexo4 pag.61: 'Cuando el COD_TIPO_PLAN_CARRERA es 1 (Plan Regular), los SEMESTRES_RECONOCIDOS deben ser 0'.", "BLOQUEANTE", MANUAL, 61, None, "Verificar SEMESTRES_RECONOCIDOS o el tipo de plan declarado.")
    if sem_rec is not None and num(sem_rec) is not None and not (0 <= num(sem_rec) <= 14):
        add(fila, row, "SEMESTRES_RECONOCIDOS", sem_rec, "Debe estar entre 0 y 14 semestres (Anexo4 pag.57/62).", "BLOQUEANTE", MANUAL, 62, None, "Corregir valor fuera de rango.")

    # 8. Duraciones
    de = num(row.get("DURACION_ESTUDIOS")); dt = num(row.get("DURACION_TITULACION")); dtot = num(row.get("DURACION_TOTAL"))
    if de is not None and not (0 < de <= 24):
        add(fila, row, "DURACION_ESTUDIOS", row.get("DURACION_ESTUDIOS"), "No puede ser 0 ni superar 24 semestres (Anexo4 pag.58/63).", "BLOQUEANTE", MANUAL, 63, None, "Corregir DURACION_ESTUDIOS.")
    if dt is not None and not (0 <= dt <= 24):
        add(fila, row, "DURACION_TITULACION", row.get("DURACION_TITULACION"), "Debe ser >=0 y <=24 semestres (Anexo4 pag.63).", "BLOQUEANTE", MANUAL, 63, None, "Corregir DURACION_TITULACION.")
    if dtot is not None and not (0 < dtot <= 24):
        add(fila, row, "DURACION_TOTAL", row.get("DURACION_TOTAL"), "No puede ser 0 ni superar 24 semestres (Anexo4 pag.63/64).", "BLOQUEANTE", MANUAL, 64, None, "Corregir DURACION_TOTAL.")
    if de is not None and dtot is not None and dtot < de:
        add(fila, row, "DURACION_TOTAL", row.get("DURACION_TOTAL"), "DURACION_TOTAL no puede ser menor a DURACION_ESTUDIOS (Anexo4 pag.64).", "BLOQUEANTE", MANUAL, 64, None, "Revisar consistencia de duraciones.")
    if de is not None and dt is not None and dtot is not None and dtot > de + dt:
        add(fila, row, "DURACION_TOTAL", row.get("DURACION_TOTAL"), "DURACION_TOTAL no puede ser mayor a DURACION_ESTUDIOS + DURACION_TITULACION (Anexo2 pag.36; Anexo4 pag.64).", "BLOQUEANTE", MANUAL, 36, 1237, "Revisar consistencia de duraciones.")

    # 9. REGIMEN=1 => DURACION_REGIMEN = DURACION_TOTAL
    dreg = num(row.get("DURACION_REGIMEN"))
    if regimen == "1" and dreg is not None and dtot is not None and dreg != dtot:
        add(fila, row, "DURACION_REGIMEN", row.get("DURACION_REGIMEN"), "Cuando REGIMEN=1 (Semestres), DURACION_REGIMEN debe ser igual a DURACION_TOTAL (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Igualar DURACION_REGIMEN a DURACION_TOTAL o corregir REGIMEN.")
    if regimen in ("2", "3", "4"):
        add(fila, row, "DURACION_REGIMEN", row.get("DURACION_REGIMEN"), f"Anexo4 (pag.59/64) indica 'No corresponde la DURACION_REGIMEN para el REGIMEN {regimen}', lo cual es ambiguo frente a la obligatoriedad general del campo (pag.58/63). No se imputa ni vacia automaticamente.", "REVISION_MANUAL", MANUAL, 59, 2018, "Confirmar con SIES si DURACION_REGIMEN debe informarse o dejarse vacio para regimenes distintos de semestral.")
    if dreg is not None and dreg <= 0:
        add(fila, row, "DURACION_REGIMEN", row.get("DURACION_REGIMEN"), "DURACION_REGIMEN debe ser mayor a 0 (Anexo4 pag.58/63).", "BLOQUEANTE", MANUAL, 63, None, "Corregir DURACION_REGIMEN.")

    # 10. Areas segun nivel
    area_actual = row.get("AREA_ACTUAL")
    area_actual_s = str(int(area_actual)) if isinstance(area_actual, (int, float)) else (str(area_actual) if area_actual is not None else None)
    area_vals = {c: row.get(c) for c in AREA_COLS}
    n_unos = sum(1 for v in area_vals.values() if num(v) == 1)
    if nivel == "1":
        target = AREA_MAP.get(area_actual_s)
        if target and num(area_vals.get(target)) != 1:
            add(fila, row, target, area_vals.get(target), f"COD_NIVEL_CARRERA=1 (TNS) y AREA_ACTUAL={area_actual_s}: el area de destino correspondiente ({target}) debe ser 1 (Anexo4 pag.60).", "BLOQUEANTE", MANUAL, 60, None, "Marcar el area de destino correspondiente en 1.")
        if n_unos > 5:
            add(fila, row, "AREA_* (destino)", n_unos, "TNS no puede tener mas de 5 areas de destino marcadas (Anexo4 pag.60).", "BLOQUEANTE", MANUAL, 60, None, "Reducir la cantidad de areas de destino marcadas.")
    elif nivel == "0":
        target = AREA_MAP.get(area_actual_s)
        if target and num(area_vals.get(target)) != 1:
            add(fila, row, target, area_vals.get(target), f"COD_NIVEL_CARRERA=0 y AREA_ACTUAL={area_actual_s}: el area de destino correspondiente ({target}) debe ser 1 (Anexo4 pag.59-60).", "BLOQUEANTE", MANUAL, 59, None, "Marcar el area de destino correspondiente en 1.")
    elif nivel == "2":
        if n_unos > 0:
            add(fila, row, "AREA_* (destino)", n_unos, "COD_NIVEL_CARRERA=2 (Profesional sin Licenciatura): todas las areas de destino deben ser 0 (Anexo4 pag.60).", "BLOQUEANTE", MANUAL, 60, None, "Vaciar/poner en 0 las areas de destino marcadas.")
    elif nivel == "3":
        educ_ok = num(area_vals.get("AREA_EDUCACION")) == 1
        others_zero = all(num(area_vals.get(c)) == 0 for c in AREA_COLS if c != "AREA_EDUCACION")
        if not (educ_ok and others_zero):
            add(fila, row, "AREA_* (destino)", str(area_vals), "COD_NIVEL_CARRERA=3: AREA_EDUCACION debe ser 1 y el resto 0 (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Ajustar areas de destino segun regla de nivel 3.")
    elif nivel == "4":
        if n_unos > 0:
            add(fila, row, "AREA_* (destino)", n_unos, "COD_NIVEL_CARRERA=4: todas las areas de destino deben ser 0 (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Vaciar/poner en 0 las areas de destino marcadas.")

    # 10b. Dominios categoricos simples (catalogo oficial)
    dominios = {
        "MODALIDAD": ("1", "2", "3"),
        "COD_JORNADA": ("1", "2", "3", "4"),
        "COD_TIPO_PLAN_CARRERA": ("1", "2", "3"),
        "REGIMEN": ("1", "2", "3", "4"),
        "COD_NIVEL_CARRERA": ("0", "1", "2", "3", "4"),
        "REQUISITO_INGRESO": ("1", "2", "3", "4", "5", "6"),
        "FORMATO_VALOR": ("1", "2"),
        "NOTAS_ENS_MEDIA": ("SI", "NO"),
        "EXPERIENCIA_LABORAL": ("SI", "NO"),
    }
    for campo_dom, dominio_ok in dominios.items():
        val = row.get(campo_dom)
        if val is None:
            continue
        val_s = str(int(val)) if isinstance(val, float) and val.is_integer() else str(val)
        if val_s not in dominio_ok:
            add(fila, row, campo_dom, val, f"Valor fuera del dominio oficial {dominio_ok} (Anexo2/Anexo4).", "BLOQUEANTE", MANUAL, "35-41 / 57-64", None, f"Corregir {campo_dom}: valor observado no pertenece al catalogo oficial.")

    acred = row.get("ACREDITACION")
    if acred is not None and str(int(acred)) not in ("1", "2", "3"):
        add(fila, row, "ACREDITACION", acred, "Fuera del dominio de la tabla de estructura (1,2,3) y del Anexo4 (1,2) (pag.37/58/63).", "BLOQUEANTE", MANUAL, 37, 1276, "Corregir ACREDITACION.")

    aact = row.get("AREA_ACTUAL")
    if aact is not None:
        aact_s = str(int(aact)) if isinstance(aact, (int, float)) else str(aact)
        if aact_s not in AREA_MAP:
            add(fila, row, "AREA_ACTUAL", aact, "Fuera del dominio oficial 1-10 (Anexo2 pag.37).", "BLOQUEANTE", MANUAL, 37, 1300, "Corregir AREA_ACTUAL.")

    for c in AREA_COLS:
        v = row.get(c)
        if v is not None and num(v) not in (0, 1):
            add(fila, row, c, v, "Cada area de destino solo admite 0 o 1 (Anexo2 pag.37-39).", "BLOQUEANTE", MANUAL, "37-39", None, f"Corregir {c} a 0 o 1.")

    licen = row.get("LICENCIA_ENS_MEDIA")
    if licen is not None and str(licen).strip().upper() != "SI":
        add(fila, row, "LICENCIA_ENS_MEDIA", licen, "Anexo4 pag.59: 'La LICENCIA_ENS_MEDIA se debe completar con SI.' (el manual no documenta explicitamente el valor NO para este campo especifico).", "REVISION_MANUAL", MANUAL, 59, None, "Confirmar con SIES si NO es un valor valido para LICENCIA_ENS_MEDIA.")

    # 11. Vigencia=1: vacantes, fecha, enlace, mail, valores
    if vigencia == "1":
        v1 = row.get("VACANTES_PRIMER_SEMESTRE"); v2 = row.get("VACANTES_SEGUNDO_SEMESTRE")
        if num(v1) == 0:
            add(fila, row, "VACANTES_PRIMER_SEMESTRE", v1, "Con VIGENCIA_CARRERA=1 las vacantes semestrales no pueden ser 0 (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Informar vacantes reales o revisar vigencia.")
        if num(v2) == 0:
            add(fila, row, "VACANTES_SEGUNDO_SEMESTRE", v2, "Con VIGENCIA_CARRERA=1 las vacantes semestrales no pueden ser 0 (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Informar vacantes reales o revisar vigencia.")
        if is_blank(row.get("FECHA_ADMISION_INICIAL")):
            add(fila, row, "FECHA_ADMISION_INICIAL", row.get("FECHA_ADMISION_INICIAL"), "Obligatoria cuando VIGENCIA_CARRERA=1 (Anexo4 pag.61).", "BLOQUEANTE", MANUAL, 61, None, "Completar la fecha de admision inicial.")
        else:
            add(fila, row, "FECHA_ADMISION_INICIAL", row.get("FECHA_ADMISION_INICIAL"), "Fecha informada, pero el rango valido para el ciclo 2027 NO esta confirmado en el instructivo (solo se encontro el rango 2025-2026, presumiblemente del ciclo anterior). No se aplica ese rango por analogia.", "REVISION_MANUAL", MANUAL, 59, 2003, "Confirmar con SIES el rango de fechas valido para FECHA_ADMISION_INICIAL del proceso 2027 antes de validar este campo.")
        enlace = row.get("ENLACE_INFO_PROGRAMA")
        if is_blank(enlace):
            add(fila, row, "ENLACE_INFO_PROGRAMA", enlace, "Obligatorio (no aplica la excepcion de vigencia=2).", "BLOQUEANTE", MANUAL, 39, 1422, "Completar enlace valido al programa.")
        elif not re.match(r"^(https?://|www\.)", str(enlace).strip(), re.IGNORECASE):
            add(fila, row, "ENLACE_INFO_PROGRAMA", enlace, "Debe iniciar con http://, https:// o www. (Anexo2 pag.39).", "ADVERTENCIA", MANUAL, 39, 1422, "Verificar formato del enlace.")
        mail = row.get("MAIL_DIFUSION_CARRERA")
        if is_blank(mail):
            add(fila, row, "MAIL_DIFUSION_CARRERA", mail, "Obligatorio cuando VIGENCIA_CARRERA=1 (Anexo4 pag.62).", "BLOQUEANTE", MANUAL, 62, None, "Completar correo de difusion.")
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", str(mail).strip()):
            add(fila, row, "MAIL_DIFUSION_CARRERA", mail, "Formato de correo invalido (Anexo4 pag.62: 'Debe ingresar un MAIL_DIFUSION_CARRERA correcto').", "BLOQUEANTE", MANUAL, 62, None, "Corregir formato de correo.")
        for campo, pag in [("EXPERIENCIA_LABORAL", 62), ("FORMATO_VALOR", 62), ("LICENCIA_ENS_MEDIA", 59), ("NOTAS_ENS_MEDIA", 59), ("COSTO_TITULACION", 62), ("VALOR_CERTIFICADO_DIPLOMA", 62), ("VALOR_MATRICULA_ANUAL", 62), ("ARANCEL_ANUAL", 61)]:
            if is_blank(row.get(campo)):
                add(fila, row, campo, row.get(campo), f"Debe completarse cuando VIGENCIA_CARRERA=1 (Anexo4).", "BLOQUEANTE", MANUAL, pag, None, f"Completar {campo}.")

    # 12. NOTAS_ENS_MEDIA <-> PROMEDIO_MIN_ENS_MEDIA
    notas = row.get("NOTAS_ENS_MEDIA")
    prom = row.get("PROMEDIO_MIN_ENS_MEDIA")
    if notas == "NO" and num(prom) not in (0, 0.0):
        add(fila, row, "PROMEDIO_MIN_ENS_MEDIA", prom, "Si NOTAS_ENS_MEDIA=NO, PROMEDIO_MIN_ENS_MEDIA debe ser 0 (Anexo2 pag.40).", "BLOQUEANTE", MANUAL, 40, 1442, "Corregir el promedio a 0.")
    elif notas == "SI" and num(prom) is not None and not (4.00 <= num(prom) <= 7.00):
        add(fila, row, "PROMEDIO_MIN_ENS_MEDIA", prom, "Debe estar entre 4.00 y 7.00 cuando NOTAS_ENS_MEDIA=SI (Anexo4 pag.58).", "BLOQUEANTE", MANUAL, 58, None, "Corregir el promedio dentro de rango.")

    # 13. Duplicado interno por llave diagnostica
    k = (norm(row.get("NOMBRE_SEDE")), norm(row.get("NOMBRE_CARRERA")), str(row.get("MODALIDAD")), str(row.get("COD_JORNADA")), str(row.get("COD_TIPO_PLAN_CARRERA")), str(row.get("COD_NIVEL_CARRERA")), norm(row.get("NOMBRE_TITULO")))
    hermanas = [f for f in key_seen[k] if f != fila]
    if hermanas:
        add(fila, row, "(fila completa)", "-", f"Llave diagnostica identica a la(s) fila(s) Excel {hermanas} dentro del mismo archivo.", "BLOQUEANTE", "perfilamiento interno", "-", None, "Confirmar con la institucion si es un registro duplicado por error de digitacion.")

    # 14. Columna no oficial SECTOR IPSS presente con dato
    if not is_blank(row.get("SECTOR IPSS")):
        add(fila, row, "SECTOR IPSS", row.get("SECTOR IPSS"), "Columna no reconocida en la estructura oficial de Etapa 2 (48 columnas, pag.35-41). No se incluye en el CSV de carga sin confirmacion.", "INFORMATIVO", MANUAL, "35-41 (ausente)", None, "No incluir en el CSV final salvo confirmacion de que es un campo oficial no documentado.")

    # 15. Marca de color roja institucional (sin interpretar su significado)
    if fila in red_by_row:
        add(fila, row, ", ".join(sorted(set(red_by_row[fila]))), "-", "Celda(s) marcadas en rojo por la institucion en el archivo original. Significado no documentado; no se asume que sea un error.", "REVISION_MANUAL", "perfilamiento interno (formato del archivo)", "-", None, "Consultar a la institucion el significado de la marca antes de dar por validos estos campos.")

    # 16. Marca de color de tema en toda la fila
    if fila in theme_rows:
        add(fila, row, "(fila completa)", "-", "Fila resaltada con color de tema aplicado a casi todas sus columnas en el archivo original. Motivo no documentado.", "REVISION_MANUAL", "perfilamiento interno (formato del archivo)", "-", None, "Consultar a la institucion el motivo del resaltado antes de validar esta fila.")

json_out = {
    "resumen": {},
    "total_filas_evaluadas": len(rows),
    "hallazgos": findings,
}
for f in findings:
    json_out["resumen"][f["severidad"]] = json_out["resumen"].get(f["severidad"], 0) + 1

json_path = os.path.join(OUT, f"VALIDACION_COMPLETA_ETAPA2_{STAMP}.json")
with open(json_path, "w", encoding="utf-8") as fh:
    json.dump(json_out, fh, ensure_ascii=False, indent=2, default=str)

tsv_path = os.path.join(OUT, f"VALIDACION_COMPLETA_ETAPA2_{STAMP}.tsv")
cols = ["fila_excel","llave_diagnostica","carrera","sede","modalidad","jornada","campo","valor_observado","regla_incumplida","fuente","pagina","linea","severidad","accion_requerida"]
with open(tsv_path, "w", encoding="utf-8") as fh:
    fh.write("\t".join(cols) + "\n")
    for d in findings:
        fh.write("\t".join(str(d.get(c,"")).replace("\t"," ").replace("\n"," ") for c in cols) + "\n")

print("Resumen severidad:", json_out["resumen"])
print("Total filas evaluadas:", len(rows))
print("Total hallazgos:", len(findings))
print("JSON:", json_path)
print("TSV:", tsv_path)
