"""Hito 3 CNED: clasificar candidatos recuperados.

Analiza archivos recuperados CNED, calcula evidencia estructural, clasifica
cada candidato y recomienda una base operativa principal preliminar.
No genera archivo final de subida CNED.
"""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover - explicit runtime guard
    raise SystemExit("ERROR: falta dependencia pandas para ejecutar Hito 3 CNED.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - explicit runtime guard
    raise SystemExit("ERROR: falta dependencia openpyxl para ejecutar Hito 3 CNED.") from exc


KEY_COLUMNS = {
    "CODIGO_UNICO",
    "CODIGO_IES",
    "CODIGO_IES_NUM",
    "CODIGO_SEDE",
    "COD_SED",
    "COD_SEDE",
    "NOMBRE_SEDE",
    "CODIGO_CARRERA",
    "COD_CARRERA",
    "COD_CAR",
    "NOMBRE_CARRERA",
    "MODALIDAD",
    "CODIGO_JORNADA",
    "COD_JORNADA",
    "JORNADA",
    "VERSION",
    "VERSION_CARRERA",
    "TIPO_CARRERA",
    "CODIGO_TIPO_CARRERA",
    "DURACION_ESTUDIOS",
    "DURACION_ESTUDIO",
    "DURACION_TITULACION",
    "DURACION_TOTAL",
    "REGIMEN",
    "AREA",
    "VIGENCIA",
    "ANIO_INICIO",
    "ANO_INICIO",
    "AÑO_INICIO",
    "ACREDITACION",
    "REQUISITO_INGRESO",
    "SEMESTRES_RECONOCIDOS",
    "NIVEL",
    "NIVEL_GLOBAL",
    "CODIGO_NIVEL_GLOBAL",
    "CODIGO_NIVEL_CARRERA",
}

STRUCTURAL_GROUPS = {
    "codigo": {"CODIGO_UNICO"},
    "carrera": {"CODIGO_CARRERA", "COD_CARRERA", "COD_CAR", "NOMBRE_CARRERA"},
    "sede": {"CODIGO_SEDE", "COD_SED", "COD_SEDE", "NOMBRE_SEDE"},
    "jornada": {"CODIGO_JORNADA", "COD_JORNADA", "JORNADA"},
    "modalidad": {"MODALIDAD"},
    "version": {"VERSION", "VERSION_CARRERA"},
}


@dataclass
class SheetProfile:
    sheet_name: str
    rows: int
    columns: int
    non_empty_rows: int
    non_empty_columns: int
    unnamed_pct: float
    headers: list[str]
    normalized_headers: list[str]
    key_columns: list[str]


@dataclass
class FileProfile:
    path: str
    name: str
    extension: str
    size_bytes: int
    sha256_short: str
    sheet_count: int
    sheet_names: list[str]
    sheets: list[SheetProfile] = field(default_factory=list)
    all_headers: list[str] = field(default_factory=list)
    all_headers_normalized: list[str] = field(default_factory=list)
    key_columns_detected: list[str] = field(default_factory=list)
    positive_signals: list[str] = field(default_factory=list)
    negative_signals: list[str] = field(default_factory=list)
    score: int = 0
    classification: str = "REQUIERE_REVISION_MANUAL"
    recommendation_rank: int = 999
    evidence_for: list[str] = field(default_factory=list)
    evidence_against: list[str] = field(default_factory=list)


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text)
        if unicodedata.category(ch) != "Mn"
    )
    for char in ["\n", "\r", "\t", "/", "-", ".", "(", ")", "[", "]", ":"]:
        text = text.replace(char, "_")
    text = "_".join(part for part in text.split() if part)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def sha256_short(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def git_status_branch(root: Path) -> str:
    result = subprocess.run(
        ["git", "status", "--short", "--branch"],
        cwd=root,
        check=True,
        text=True,
        capture_output=True,
    )
    return result.stdout.strip()


def read_table_file(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    return pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False)


def profile_dataframe(sheet_name: str, df: pd.DataFrame) -> SheetProfile:
    headers = ["" if col is None else str(col) for col in df.columns]
    normalized = [normalize_header(col) for col in headers]
    unnamed = sum(
        1 for raw, norm in zip(headers, normalized)
        if not norm or norm.startswith("UNNAMED") or raw.startswith("Unnamed:")
    )
    non_empty_rows = int(df.replace("", pd.NA).dropna(how="all").shape[0])
    non_empty_columns = int(df.replace("", pd.NA).dropna(how="all", axis=1).shape[1])
    key_columns = sorted(set(normalized) & KEY_COLUMNS)
    return SheetProfile(
        sheet_name=sheet_name,
        rows=int(df.shape[0]),
        columns=int(df.shape[1]),
        non_empty_rows=non_empty_rows,
        non_empty_columns=non_empty_columns,
        unnamed_pct=round((unnamed / len(headers) * 100) if headers else 0, 2),
        headers=headers,
        normalized_headers=normalized,
        key_columns=key_columns,
    )


def profile_file(path: Path) -> FileProfile:
    extension = path.suffix.lower()
    profile = FileProfile(
        path=str(path),
        name=path.name,
        extension=extension,
        size_bytes=path.stat().st_size,
        sha256_short=sha256_short(path),
        sheet_count=0,
        sheet_names=[],
    )

    if extension == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        profile.sheet_names = list(workbook.sheetnames)
        profile.sheet_count = len(profile.sheet_names)
        for sheet in profile.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet, dtype=str, keep_default_na=False)
            profile.sheets.append(profile_dataframe(sheet, df))
    elif extension in {".csv", ".tsv"}:
        df = read_table_file(path)
        profile.sheet_names = [path.stem]
        profile.sheet_count = 1
        profile.sheets.append(profile_dataframe(path.stem, df))
    else:
        profile.negative_signals.append(f"Extension no soportada: {extension}")

    all_headers: list[str] = []
    all_normalized: list[str] = []
    for sheet_profile in profile.sheets:
        all_headers.extend(sheet_profile.headers)
        all_normalized.extend(sheet_profile.normalized_headers)
    profile.all_headers = sorted(set(all_headers))
    profile.all_headers_normalized = sorted(set(all_normalized))
    profile.key_columns_detected = sorted(set(all_normalized) & KEY_COLUMNS)
    score_and_classify(profile)
    return profile


def has_all_structural_groups(headers: set[str]) -> bool:
    return all(headers & options for options in STRUCTURAL_GROUPS.values())


def add_score(profile: FileProfile, points: int, signal: str, positive: bool = True) -> None:
    profile.score += points
    if positive:
        profile.positive_signals.append(f"{signal} ({points:+d})")
    else:
        profile.negative_signals.append(f"{signal} ({points:+d})")


def score_and_classify(profile: FileProfile) -> None:
    name = normalize_header(profile.name)
    headers = set(profile.all_headers_normalized)
    max_rows = max((sheet.non_empty_rows for sheet in profile.sheets), default=0)
    max_cols = max((sheet.non_empty_columns for sheet in profile.sheets), default=0)
    avg_unnamed = (
        sum(sheet.unnamed_pct for sheet in profile.sheets) / len(profile.sheets)
        if profile.sheets else 100
    )

    for marker, points in [
        ("SANEADO", 30),
        ("FINAL", 25),
        ("OPERATIVO", 20),
        ("LISTO_CARGA", 20),
        ("CON_COD_CARRERA", 15),
        ("CREACION_PROGRAMAS", 12),
    ]:
        if marker in name:
            add_score(profile, points, f"Nombre contiene {marker}")

    if has_all_structural_groups(headers):
        add_score(profile, 30, "Tiene codigo/carrera/sede/jornada/modalidad/version")
    if max_rows > 20 and max_cols > 10:
        add_score(profile, 15, "Tiene mas de 20 filas y mas de 10 columnas")
    if any(sheet.non_empty_rows > 0 and sheet.non_empty_columns > 1 for sheet in profile.sheets):
        add_score(profile, 15, "Tiene estructura tabular clara")
    if avg_unnamed <= 10:
        add_score(profile, 10, "Tiene pocas columnas vacias o sin nombre")
    if "CODIGO_UNICO" in headers:
        add_score(profile, 20, "Tiene CODIGO_UNICO o equivalente")
    if headers & {"DURACION_ESTUDIOS", "DURACION_ESTUDIO", "DURACION_TITULACION", "DURACION_TOTAL", "VIGENCIA"}:
        add_score(profile, 10, "Tiene campos de duracion o vigencia")

    for marker, points in [
        ("PRUEBA", -30),
        ("PARCIAL", -20),
        ("SIN_DECISION", -25),
        ("TABLA_EJECUTIVA", -15),
        ("AUDITORIA", -20),
        ("CONFLICTOS", -25),
        ("DUDOSOS", -20),
        ("DESCARTADOS", -80),
    ]:
        if marker in name:
            add_score(profile, points, f"Nombre contiene {marker}", positive=False)

    if max_rows < 15:
        add_score(profile, -15, "Tiene menos de 15 filas", positive=False)
    if profile.sheet_count == 1 and max_rows <= 15 and max_cols <= 10:
        add_score(profile, -20, "Parece tabla resumen sin estructura de carga", positive=False)
    if avg_unnamed > 30:
        add_score(profile, -15, "Tiene muchas columnas vacias o sin nombre", positive=False)

    strong_negative = any(marker in name for marker in ["PRUEBA", "PARCIAL", "SIN_DECISION", "AUDITORIA", "CONFLICTOS", "DESCARTADOS"])
    if "PRUEBA_SUBIDA" in name or "10_CARRERAS" in name:
        classification = "PRUEBA_SUBIDA"
    elif "PARCIAL" in name or "SIN_DECISION" in name:
        classification = "CARGA_PARCIAL_PENDIENTE"
    elif "TABLA_EJECUTIVA_DECISION" in name:
        classification = "INSUMO_DECISION_INSTITUCIONAL"
    elif "GUIA" in name:
        classification = "GUIA_OPERATIVA"
    elif "AUDITORIA" in name:
        classification = "AUDITORIA_REFERENCIA"
    elif profile.score >= 80 and not strong_negative:
        classification = "BASE_OPERATIVA_PRINCIPAL_PRELIMINAR"
    elif profile.score >= 55:
        classification = "BASE_OPERATIVA_SECUNDARIA"
    elif "DESCARTADOS" in name:
        classification = "DESCARTABLE"
    else:
        classification = "REQUIERE_REVISION_MANUAL"
    profile.classification = classification
    profile.evidence_for = list(profile.positive_signals)
    profile.evidence_against = list(profile.negative_signals)


def compare_profiles(a: FileProfile, b: FileProfile) -> dict[str, Any]:
    a_headers = set(a.all_headers_normalized)
    b_headers = set(b.all_headers_normalized)
    union = a_headers | b_headers
    intersection = a_headers & b_headers
    similarity = round(len(intersection) / len(union) * 100, 2) if union else 0.0
    return {
        "archivo_a": a.name,
        "archivo_b": b.name,
        "columnas_a": len(a_headers),
        "columnas_b": len(b_headers),
        "interseccion": len(intersection),
        "solo_a": sorted(a_headers - b_headers),
        "solo_b": sorted(b_headers - a_headers),
        "similitud_pct": similarity,
        "posible_derivado": similarity >= 70,
    }


def profile_to_row(profile: FileProfile) -> dict[str, Any]:
    max_rows = max((sheet.non_empty_rows for sheet in profile.sheets), default=0)
    max_cols = max((sheet.non_empty_columns for sheet in profile.sheets), default=0)
    return {
        "archivo": profile.name,
        "extension": profile.extension,
        "tamano_bytes": profile.size_bytes,
        "sha256_short": profile.sha256_short,
        "hojas": profile.sheet_count,
        "nombres_hojas": " | ".join(profile.sheet_names),
        "max_filas_no_vacias": max_rows,
        "max_columnas_no_vacias": max_cols,
        "columnas_clave": " | ".join(profile.key_columns_detected),
        "score": profile.score,
        "clasificacion": profile.classification,
        "senales_positivas": " || ".join(profile.positive_signals),
        "senales_negativas": " || ".join(profile.negative_signals),
    }


def to_jsonable(profile: FileProfile) -> dict[str, Any]:
    return {
        "path": profile.path,
        "name": profile.name,
        "extension": profile.extension,
        "size_bytes": profile.size_bytes,
        "sha256_short": profile.sha256_short,
        "sheet_count": profile.sheet_count,
        "sheet_names": profile.sheet_names,
        "sheets": [sheet.__dict__ for sheet in profile.sheets],
        "all_headers_normalized": profile.all_headers_normalized,
        "key_columns_detected": profile.key_columns_detected,
        "positive_signals": profile.positive_signals,
        "negative_signals": profile.negative_signals,
        "score": profile.score,
        "classification": profile.classification,
    }


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0].keys()) if rows else ["sin_datos"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def markdown_table(rows: list[dict[str, Any]], columns: list[str]) -> list[str]:
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    for row in rows:
        values = [str(row.get(col, "")).replace("|", "/") for col in columns]
        lines.append("| " + " | ".join(values) + " |")
    return lines


def main() -> int:
    root = repo_root()
    input_dir = root / "cned" / "data" / "input" / "recuperacion_controlada"
    report_dir = root / "cned" / "resultados" / "reportes"
    audit_dir = root / "cned" / "resultados" / "auditorias"
    processed_dir = root / "cned" / "data" / "processed"
    for directory in [report_dir, audit_dir, processed_dir]:
        directory.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    md_path = report_dir / f"CNED_HITO3_CLASIFICACION_CANDIDATOS_{timestamp}.md"
    csv_path = audit_dir / f"CNED_HITO3_CLASIFICACION_CANDIDATOS_{timestamp}.csv"
    json_path = audit_dir / f"CNED_HITO3_CLASIFICACION_CANDIDATOS_{timestamp}.json"

    candidates = sorted(
        path for path in input_dir.iterdir()
        if path.is_file() and path.suffix.lower() in {".xlsx", ".csv", ".tsv"}
    )
    if not candidates:
        raise SystemExit(f"ERROR: no hay candidatos en {input_dir}")

    profiles = [profile_file(path) for path in candidates]
    ranked = sorted(profiles, key=lambda item: item.score, reverse=True)
    for idx, profile in enumerate(ranked, start=1):
        profile.recommendation_rank = idx

    principal = next((p for p in ranked if p.classification == "BASE_OPERATIVA_PRINCIPAL_PRELIMINAR"), None)
    if principal is None:
        principal = ranked[0]
        if principal.classification not in {"PRUEBA_SUBIDA", "CARGA_PARCIAL_PENDIENTE", "INSUMO_DECISION_INSTITUCIONAL"}:
            principal.classification = "BASE_OPERATIVA_SECUNDARIA"

    comparisons: dict[str, dict[str, Any]] = {}
    by_prefix = {profile.name[:2]: profile for profile in profiles}
    for label, left, right in [
        ("comparacion_01_vs_04", "01", "04"),
        ("comparacion_01_vs_02", "01", "02"),
        ("comparacion_01_vs_07", "01", "07"),
    ]:
        if left in by_prefix and right in by_prefix:
            comparisons[label] = compare_profiles(by_prefix[left], by_prefix[right])

    copied_path = None
    principal_path = Path(principal.path)
    if principal_path.suffix.lower() == ".xlsx":
        copied_path = processed_dir / f"CNED_BASE_OPERATIVA_RECOMENDADA_PRELIMINAR_{timestamp}.xlsx"
        shutil.copy2(principal_path, copied_path)

    rows = [profile_to_row(profile) for profile in ranked]
    write_csv(csv_path, rows)
    json_payload = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "git_status_inicial": git_status_branch(root),
        "input_dir": str(input_dir),
        "recomendacion_preliminar": principal.name,
        "score_recomendado": principal.score,
        "clasificacion_recomendada": principal.classification,
        "copia_preliminar": str(copied_path) if copied_path else None,
        "ranking": rows,
        "profiles": [to_jsonable(profile) for profile in ranked],
        "comparaciones": comparisons,
        "advertencia": "No se genera archivo final CNED hasta confirmar instructivo/estructura oficial.",
    }
    json_path.write_text(json.dumps(json_payload, ensure_ascii=False, indent=2), encoding="utf-8")

    class_counts = Counter(profile.classification for profile in profiles)
    md_lines = [
        "# CNED Hito 3 - Clasificacion de candidatos recuperados",
        "",
        f"- Fecha: {datetime.now().isoformat(timespec='seconds')}",
        f"- Directorio analizado: `{input_dir.relative_to(root)}`",
        f"- Archivos analizados: {len(profiles)}",
        f"- Estado git inicial:",
        "",
        "```text",
        git_status_branch(root),
        "```",
        "",
        "## Ranking",
        "",
        *markdown_table(
            rows,
            ["archivo", "score", "clasificacion", "max_filas_no_vacias", "max_columnas_no_vacias", "columnas_clave"],
        ),
        "",
        "## Recomendacion preliminar",
        "",
        f"- Base operativa principal preliminar: `{principal.name}`",
        f"- Puntaje: {principal.score}",
        f"- Clasificacion: {principal.classification}",
        f"- Copia preliminar: `{copied_path.relative_to(root) if copied_path else 'NO_GENERADA'}`",
        "",
        "### Evidencia a favor",
        "",
        *(f"- {signal}" for signal in principal.evidence_for),
        "",
        "### Evidencia en contra",
        "",
        *(f"- {signal}" for signal in (principal.evidence_against or ["Sin senales negativas relevantes."])),
        "",
        "## Clasificaciones",
        "",
    ]
    md_lines.extend(f"- {classification}: {count}" for classification, count in sorted(class_counts.items()))
    md_lines.extend(["", "## Comparaciones estructurales", ""])
    for label, comparison in comparisons.items():
        md_lines.extend([
            f"### {label}",
            "",
            f"- Archivo A: `{comparison['archivo_a']}`",
            f"- Archivo B: `{comparison['archivo_b']}`",
            f"- Similitud de encabezados: {comparison['similitud_pct']}%",
            f"- Interseccion columnas: {comparison['interseccion']}",
            f"- Posible derivado: {'SI' if comparison['posible_derivado'] else 'NO'}",
            f"- Solo A: {', '.join(comparison['solo_a'][:30]) or 'Sin diferencias'}",
            f"- Solo B: {', '.join(comparison['solo_b'][:30]) or 'Sin diferencias'}",
            "",
        ])
    md_lines.extend([
        "## Interpretacion ejecutiva",
        "",
        "- `01` concentra senales fuertes de base saneada/final/operativa si mantiene estructura tabular y columnas clave.",
        "- `02` y `03` se consideran pruebas de subida si contienen `10_CARRERAS` o `PRUEBA_SUBIDA`.",
        "- `07` y `10` se consideran cargas parciales pendientes si contienen `PARCIAL` o `SIN_DECISION`.",
        "- `08` y `09` se consideran insumos de decision institucional, no bases operativas.",
        "",
        "## Proximos pasos",
        "",
        "1. Confirmar instructivo o estructura oficial CNED.",
        "2. Definir columnas finales y reglas de transformacion autorizadas.",
        "3. Construir auditoria de completitud sobre la base recomendada preliminar.",
        "4. Generar archivo de subida solo cuando no existan pendientes metodologicos.",
        "",
        "> Advertencia: No se genera archivo final CNED hasta confirmar instructivo/estructura oficial.",
        "",
        "## Dictamen",
        "",
        "DICTAMEN_FINAL: HITO3_CLASIFICACION_CANDIDATOS_CNED_GENERADA",
    ])
    md_path.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    print("======================================================================")
    print("HITO 3 CNED - CLASIFICACION CANDIDATOS")
    print("======================================================================")
    print(f"Archivos analizados: {len(profiles)}")
    print(f"Base recomendada preliminar: {principal.name}")
    print(f"Puntaje: {principal.score}")
    print(f"Clasificacion: {principal.classification}")
    print(f"Reporte Markdown: {md_path}")
    print(f"Auditoria CSV: {csv_path}")
    print(f"Auditoria JSON: {json_path}")
    print(f"Copia preliminar: {copied_path if copied_path else 'NO_GENERADA'}")
    print("Advertencia: No se genera archivo final CNED hasta confirmar instructivo/estructura oficial.")
    print("DICTAMEN_FINAL: HITO3_CLASIFICACION_CANDIDATOS_CNED_GENERADA")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
