"""Hito 7 CNED: plantilla de resolucion institucional.

Construye una plantilla editable para resolver campos oficiales bloqueantes y
clasificar campos de formulario. No aplica cambios automaticos ni genera carga
final.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta pandas para ejecutar Hito 7 CNED.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta openpyxl para ejecutar Hito 7 CNED.") from exc


HITO6_XLSX = Path("cned/resultados/archivos_subida/CNED_MATRIZ_RESOLUCION_BLOQUEOS_HITO6_20260605_140333.xlsx")
HITO5_XLSX = Path("cned/resultados/archivos_subida/CNED_BASE_SANEADA_PRELIMINAR_HITO5_20260605_135231.xlsx")

OFFICIAL_FIELDS = {
    "Régimen": ["REGIMEN", "REGIMEN", "Régimen"],
    "VIGENCIA_CARRERA": ["VIGENCIA", "VIGENCIA_CARRERA"],
    "ACREDITACION": ["ACREDITACION", "ACREDITACIÓN"],
    "ANIO_INICIO": ["ANIO_INICIO", "AÑO_INICIO", "ANO_INICIO", "Año de inicio de Actividades"],
    "AREA_ACTUAL": ["AREA_ACTUAL", "AREA", "ÁREA", "Área del Conocimiento"],
    "DURACION_REGIMEN": ["DURACION_REGIMEN", "DURACIÓN_RÉGIMEN", "Duración del programa en Semestres"],
    "NOMBRE_TITULO": ["NOMBRE_TITULO", "NOMBRE_TÍTULO", "TITULO", "TÍTULO", "Título que otorga el programa"],
    "REQUISITO_INGRESO": ["REQUISITO_INGRESO", "REQUISITO_DE_INGRESO", "TEXTO_REQUISITO_INGRESO"],
    "SEMESTRES_RECONOCIDOS": ["SEMESTRES_RECONOCIDOS"],
    "REGIMEN": ["REGIMEN", "Régimen"],
}

TECH_PATTERNS = [
    "REVISAR",
    "CONFLICTO",
    "DECISION",
    "PENDIENTE",
    "VALIDAR",
    "VERIFICAR",
    "DUDOSO",
    "SIN_FUENTE",
    "SIN DECISION",
    "SIN_DECISION",
    "NO_DEFINIDO",
    "POR_DEFINIR",
]


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def normalize(value: Any) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = text.strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text)
        if unicodedata.category(ch) != "Mn"
    )
    for char in [" ", ".", "/", "\\", "-", "(", ")", "[", "]", "\n", "\r", "\t", ":"]:
        text = text.replace(char, "_")
    text = re.sub(r"[^A-Z0-9_]+", "_", text)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def norm_value(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            return text
    return text


def contains_mark(value: Any) -> bool:
    text = normalize(value)
    return any(pattern in text for pattern in TECH_PATTERNS)


def run_guardrail(root: Path) -> tuple[bool, str]:
    py = root / ".venv" / "bin" / "python"
    cmd = [str(py) if py.exists() else "python", "cned/scripts/check_no_tocar_mu2026.py"]
    result = subprocess.run(cmd, cwd=root, text=True, capture_output=True)
    return result.returncode == 0, (result.stdout + result.stderr).strip()


def git_status(root: Path) -> str:
    result = subprocess.run(["git", "status", "--short", "--branch"], cwd=root, text=True, capture_output=True, check=True)
    return result.stdout.strip()


def colmap(df: pd.DataFrame) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for col in df.columns:
        mapping.setdefault(normalize(col), str(col))
    return mapping


def get_existing_columns(df: pd.DataFrame, candidates: list[str]) -> list[str]:
    mapping = colmap(df)
    out = []
    for candidate in candidates:
        col = mapping.get(normalize(candidate))
        if col and col not in out:
            out.append(col)
    return out


def unique_clean_values(df: pd.DataFrame, col: str) -> list[str]:
    values = [norm_value(v) for v in df[col].tolist()]
    return sorted({v for v in values if v and not contains_mark(v)})


def unique_marked_values(df: pd.DataFrame, col: str) -> list[str]:
    values = [norm_value(v) for v in df[col].tolist()]
    return sorted({v for v in values if v and contains_mark(v)})


def proposal_for_field(field: str, candidates: list[str], base: pd.DataFrame, official_summary: pd.DataFrame) -> dict[str, Any]:
    existing = get_existing_columns(base, candidates)
    clean_sources = []
    marked_sources = []
    value_signatures: dict[str, list[str]] = {}
    for col in existing:
        clean_values = unique_clean_values(base, col)
        marked_values = unique_marked_values(base, col)
        if clean_values and not marked_values:
            clean_sources.append(col)
            value_signatures[col] = clean_values
        elif clean_values:
            marked_sources.append(col)
            value_signatures[col] = clean_values[:20]
        elif marked_values:
            marked_sources.append(col)
            value_signatures[col] = marked_values[:20]

    if clean_sources:
        signatures = {tuple(vals) for vals in value_signatures.values() if vals}
        if len(signatures) == 1:
            propuesta = "RESUELTO_POR_EQUIVALENCIA_INTERNA_PROPUESTA"
            requiere = "VALIDAR_EQUIVALENCIA_INTERNA"
        else:
            propuesta = "EQUIVALENCIAS_INTERNAS_CON_VALORES_DISTINTOS"
            requiere = "DECISION_MANUAL"
    elif marked_sources:
        propuesta = "SIN_EQUIVALENCIA_LIMPIA"
        requiere = "FUENTE_INSTITUCIONAL"
    else:
        propuesta = "NO_EXISTE_COLUMNA_EQUIVALENTE"
        requiere = "FUENTE_INSTITUCIONAL"

    affected = ""
    total_marks = ""
    if not official_summary.empty and "campo" in official_summary.columns:
        match = official_summary[official_summary["campo"].astype(str) == field]
        if not match.empty:
            affected = match.iloc[0].get("total_registros_afectados", "")
            total_marks = match.iloc[0].get("total_marcas", "")

    return {
        "CAMPO_OFICIAL": field,
        "EQUIVALENCIAS_BUSCADAS": " | ".join(candidates),
        "COLUMNAS_EXISTENTES": " | ".join(existing),
        "COLUMNAS_EQUIVALENTES_LIMPIAS": " | ".join(clean_sources),
        "COLUMNAS_CON_MARCAS": " | ".join(marked_sources),
        "VALORES_LIMPIOS_EJEMPLO": json.dumps(value_signatures, ensure_ascii=False)[:1000],
        "TOTAL_REGISTROS_AFECTADOS": affected,
        "TOTAL_MARCAS": total_marks,
        "PROPUESTA_RESOLUCION": propuesta,
        "FUENTE_REQUERIDA": requiere,
        "APLICAR_AUTOMATICAMENTE": "NO",
        "OBSERVACION": "Propuesta de resolucion; no se aplica reemplazo automatico.",
    }


def classify_form_field(row: pd.Series) -> dict[str, Any]:
    field = str(row.get("campo", ""))
    norm = normalize(field)
    if any(token in norm for token in ["ARANCEL", "VALOR", "MATRICULA", "COSTO", "CERTIFICADO"]):
        clas = "REQUIERE_DECISION_INSTITUCIONAL"
        source = "FINANZAS_ARANCELES"
    elif any(token in norm for token in ["MALLA", "PERFIL", "REQUISITO", "INGRESO", "AREA", "SUB_AREA", "DEPENDENCIA"]):
        clas = "REQUIERE_DECISION_INSTITUCIONAL"
        source = "AREA_ACADEMICA_INSTRUCTIVO_CNED"
    elif any(token in norm for token in ["OBSERVACION"]):
        clas = "SOLO_TRAZABILIDAD_O_FORMULARIO"
        source = "TRAZABILIDAD"
    else:
        clas = "SOLO_TRAZABILIDAD_O_FORMULARIO"
        source = "INSTRUCTIVO_CNED"
    return {
        "campo": field,
        "tipo_campo": row.get("tipo_campo", "FORMULARIO_CNED_DESCONOCIDO"),
        "total_marcas": row.get("total_marcas", ""),
        "total_registros_afectados": row.get("total_registros_afectados", ""),
        "clasificacion_hito7": clas,
        "fuente_o_decision_requerida": source,
        "decision_institucional": "",
        "observacion": "No excluir ni completar sin instructivo CNED o decision institucional.",
    }


def build_decision_matrix(official: pd.DataFrame, form: pd.DataFrame, aux: pd.DataFrame, proposals: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, row in official.iterrows():
        field = str(row.get("campo", ""))
        prop = proposals[proposals["CAMPO_OFICIAL"] == field]
        propuesta = prop.iloc[0].get("PROPUESTA_RESOLUCION", "COMPLETAR_DESDE_FUENTE_OFICIAL") if not prop.empty else "COMPLETAR_DESDE_FUENTE_OFICIAL"
        fuente = prop.iloc[0].get("FUENTE_REQUERIDA", "FUENTE_INSTITUCIONAL") if not prop.empty else "FUENTE_INSTITUCIONAL"
        rows.append({
            "CAMPO": field,
            "TIPO_CAMPO": "OFICIAL_PROBABLE",
            "TOTAL_REGISTROS_AFECTADOS": row.get("total_registros_afectados", ""),
            "TOTAL_MARCAS": row.get("total_marcas", ""),
            "VALORES_ACTUALES_EJEMPLO": "",
            "PROPUESTA_ACCION": propuesta,
            "FUENTE_REQUERIDA": fuente,
            "RESPONSABLE": "",
            "DECISION_INSTITUCIONAL": "",
            "VALOR_CORREGIDO_SI_APLICA": "",
            "OBSERVACION": "Completar solo con fuente trazable; no aplicar automaticamente.",
            "ESTADO_DECISION": "PENDIENTE",
        })
    for _, row in form.iterrows():
        rows.append({
            "CAMPO": row.get("campo", ""),
            "TIPO_CAMPO": "FORMULARIO_CNED_DESCONOCIDO",
            "TOTAL_REGISTROS_AFECTADOS": row.get("total_registros_afectados", ""),
            "TOTAL_MARCAS": row.get("total_marcas", ""),
            "VALORES_ACTUALES_EJEMPLO": "",
            "PROPUESTA_ACCION": row.get("clasificacion_hito7", ""),
            "FUENTE_REQUERIDA": row.get("fuente_o_decision_requerida", ""),
            "RESPONSABLE": "",
            "DECISION_INSTITUCIONAL": "",
            "VALOR_CORREGIDO_SI_APLICA": "",
            "OBSERVACION": row.get("observacion", ""),
            "ESTADO_DECISION": "PENDIENTE",
        })
    for _, row in aux.iterrows():
        rows.append({
            "CAMPO": row.get("campo", ""),
            "TIPO_CAMPO": "AUXILIAR_TECNICA",
            "TOTAL_REGISTROS_AFECTADOS": row.get("total_registros_afectados", ""),
            "TOTAL_MARCAS": row.get("total_marcas", ""),
            "VALORES_ACTUALES_EJEMPLO": "",
            "PROPUESTA_ACCION": "EXCLUIR_DE_CARGA",
            "FUENTE_REQUERIDA": "TRAZABILIDAD_INTERNA",
            "RESPONSABLE": "",
            "DECISION_INSTITUCIONAL": "",
            "VALOR_CORREGIDO_SI_APLICA": "",
            "OBSERVACION": "Separar de archivo de carga y conservar en auditoria.",
            "ESTADO_DECISION": "PENDIENTE",
        })
    return pd.DataFrame(rows)


def build_candidate(base: pd.DataFrame, official_fields: list[str]) -> pd.DataFrame:
    aux_cols = [col for col in base.columns if normalize(col).startswith("CNED_") or "MATCH" in normalize(col)]
    candidate = base.drop(columns=aux_cols, errors="ignore").copy()
    candidate["HITO7_REQUIERE_DECISION"] = "SI"
    candidate["HITO7_CAMPOS_PENDIENTES"] = " | ".join(official_fields)
    candidate["HITO7_ACCION_RECOMENDADA"] = "Resolver campos oficiales y decisiones de formulario; luego reejecutar Hito 5."
    candidate["HITO7_LISTO_PARA_REEJECUTAR_HITO5"] = "NO"
    return candidate


def md_table(df: pd.DataFrame, columns: list[str], max_rows: int = 50) -> list[str]:
    if df.empty:
        return ["_Sin datos._"]
    cols = [c for c in columns if c in df.columns]
    subset = df.loc[:, cols].head(max_rows)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
    for _, row in subset.iterrows():
        lines.append("| " + " | ".join(str(row[c]).replace("|", "/") for c in cols) + " |")
    return lines


def safe(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy().astype(str)


def write_csv(path: Path, df: pd.DataFrame) -> None:
    df.to_csv(path, index=False, quoting=csv.QUOTE_MINIMAL)


def main() -> int:
    root = repo_root()
    guard_ok, guard_output = run_guardrail(root)
    if not guard_ok:
        print(guard_output)
        return 2
    hito6 = root / HITO6_XLSX
    hito5 = root / HITO5_XLSX
    if not hito6.exists():
        print(f"ERROR: no existe matriz Hito 6: {hito6}")
        return 1
    if not hito5.exists():
        print(f"ERROR: no existe base Hito 5: {hito5}")
        return 1

    official = pd.read_excel(hito6, sheet_name="CAMPOS_OFICIALES", dtype=str, keep_default_na=False)
    form_raw = pd.read_excel(hito6, sheet_name="CAMPOS_FORMULARIO", dtype=str, keep_default_na=False)
    aux = pd.read_excel(hito6, sheet_name="CAMPOS_AUXILIARES", dtype=str, keep_default_na=False)
    base = pd.read_excel(hito5, sheet_name="CANDIDATO_CARGA_LIMPIO", dtype=str, keep_default_na=False)
    base_ref = base.head(200).copy()

    proposal_rows = [
        proposal_for_field(field, equivalents, base, official)
        for field, equivalents in OFFICIAL_FIELDS.items()
    ]
    proposals = pd.DataFrame(proposal_rows)
    form = pd.DataFrame([classify_form_field(row) for _, row in form_raw.iterrows()])
    decision_matrix = build_decision_matrix(official, form, aux, proposals)
    candidate = build_candidate(base, official["campo"].astype(str).tolist())

    internal_possible = int(proposals["PROPUESTA_RESOLUCION"].str.contains("EQUIVALENCIA_INTERNA", na=False).sum())
    require_source = int((proposals["FUENTE_REQUERIDA"] == "FUENTE_INSTITUCIONAL").sum())
    semaforo = "ROJO" if len(official) else "AMARILLO"

    plan = pd.DataFrame([
        {"orden": 1, "tarea": "Completar matriz institucional", "detalle": "Asignar responsable, decision institucional y estado para cada campo."},
        {"orden": 2, "tarea": "Validar equivalencias internas", "detalle": "Revisar propuestas RESUELTO_POR_EQUIVALENCIA_INTERNA_PROPUESTA antes de aplicar."},
        {"orden": 3, "tarea": "Levantar fuentes oficiales", "detalle": "Completar campos sin equivalencia limpia desde Registro Curricular, Oferta Academica o Finanzas."},
        {"orden": 4, "tarea": "Clasificar campos formulario", "detalle": "Definir requerido, excluible, solo trazabilidad o no aplica segun instructivo CNED."},
        {"orden": 5, "tarea": "Aplicar resoluciones trazables", "detalle": "Crear una corrida posterior que aplique solo decisiones RESUELTAS."},
        {"orden": 6, "tarea": "Reejecutar Hito 5", "detalle": "Validar que desaparezcan bloqueos oficiales antes de candidato de subida."},
    ])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = root / "cned" / "resultados" / "archivos_subida" / f"CNED_HITO7_PLANTILLA_RESOLUCION_INSTITUCIONAL_{timestamp}.xlsx"
    out_md = root / "cned" / "resultados" / "reportes" / f"CNED_HITO7_RESOLUCION_CAMPOS_OFICIALES_{timestamp}.md"
    out_official = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO7_CAMPOS_OFICIALES_{timestamp}.csv"
    out_form = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO7_CAMPOS_FORMULARIO_{timestamp}.csv"
    out_json = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO7_RESUMEN_{timestamp}.json"
    for path in [out_xlsx, out_md, out_official, out_form, out_json]:
        path.parent.mkdir(parents=True, exist_ok=True)

    resumen = pd.DataFrame([
        {"campo": "fecha_ejecucion", "valor": datetime.now().isoformat(timespec="seconds")},
        {"campo": "campos_oficiales_bloqueantes", "valor": len(official)},
        {"campo": "equivalencia_interna_posible", "valor": internal_possible},
        {"campo": "requieren_fuente_institucional", "valor": require_source},
        {"campo": "campos_formulario_decision", "valor": len(form)},
        {"campo": "campos_auxiliares_excluibles", "valor": len(aux)},
        {"campo": "semaforo", "valor": semaforo},
        {"campo": "advertencia", "valor": "Plantilla institucional; no es archivo final ni aplica cambios automaticos."},
    ])

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        safe(resumen).to_excel(writer, sheet_name="RESUMEN_EJECUTIVO", index=False)
        safe(official).to_excel(writer, sheet_name="CAMPOS_OFICIALES", index=False)
        safe(proposals).to_excel(writer, sheet_name="PROPUESTA_EQUIVALENCIAS", index=False)
        safe(form).to_excel(writer, sheet_name="CAMPOS_FORMULARIO", index=False)
        safe(aux).to_excel(writer, sheet_name="CAMPOS_AUXILIARES", index=False)
        safe(decision_matrix).to_excel(writer, sheet_name="MATRIZ_DECISION", index=False)
        safe(base_ref).to_excel(writer, sheet_name="BASE_REFERENCIA", index=False)
        safe(plan).to_excel(writer, sheet_name="PLAN_ACCION_HITO7", index=False)
        safe(candidate).to_excel(writer, sheet_name="CANDIDATO_AUX_SEPARADAS", index=False)

    write_csv(out_official, proposals)
    write_csv(out_form, form)

    payload = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "git_status_inicial": git_status(root),
        "matriz_hito6": str(HITO6_XLSX),
        "base_hito5": str(HITO5_XLSX),
        "campos_oficiales_bloqueantes": official["campo"].astype(str).tolist(),
        "equivalencia_interna_posible": internal_possible,
        "requieren_fuente_institucional": require_source,
        "campos_formulario_decision": len(form),
        "campos_auxiliares_excluibles": aux["campo"].astype(str).tolist() if "campo" in aux.columns else [],
        "semaforo": semaforo,
        "archivos_generados": {
            "excel": str(out_xlsx),
            "markdown": str(out_md),
            "csv_oficiales": str(out_official),
            "csv_formulario": str(out_form),
            "json": str(out_json),
        },
        "guardrail_mu2026": guard_output,
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    md_lines = [
        "# CNED Hito 7 - Resolucion de campos oficiales",
        "",
        f"- Fecha: {datetime.now().isoformat(timespec='seconds')}",
        f"- Campos oficiales bloqueantes: {len(official)}",
        f"- Equivalencias internas posibles: {internal_possible}",
        f"- Requieren fuente institucional: {require_source}",
        f"- Campos formulario a decidir: {len(form)}",
        f"- Campos auxiliares excluibles: {len(aux)}",
        f"- Semaforo: **{semaforo}**",
        "",
        "## Campos oficiales bloqueantes",
        "",
        *md_table(official, ["campo", "total_marcas", "total_registros_afectados", "accion_recomendada", "fuente_sugerida"], 50),
        "",
        "## Propuesta de equivalencias internas",
        "",
        *md_table(proposals, ["CAMPO_OFICIAL", "COLUMNAS_EXISTENTES", "COLUMNAS_EQUIVALENTES_LIMPIAS", "PROPUESTA_RESOLUCION", "FUENTE_REQUERIDA"], 50),
        "",
        "## Campos formulario que requieren decision",
        "",
        *md_table(form, ["campo", "total_marcas", "total_registros_afectados", "clasificacion_hito7", "fuente_o_decision_requerida"], 80),
        "",
        "## Campos auxiliares excluibles",
        "",
        *md_table(aux, ["campo", "total_marcas", "accion_recomendada", "se_puede_excluir"], 50),
        "",
        "## Plan de accion",
        "",
        *md_table(plan, ["orden", "tarea", "detalle"], 20),
        "",
        "## Condicion para volver a Hito 5",
        "",
        "Completar la matriz institucional, validar equivalencias internas y fuentes oficiales, aplicar resoluciones trazables en una corrida posterior y volver a ejecutar Hito 5.",
        "",
        "## Guardrail MU2026/PES_READY",
        "",
        "```text",
        guard_output,
        "```",
        "",
        "> Advertencia: no es archivo final CNED y no aplica cambios automaticos.",
        "",
        "## Dictamen",
        "",
        "DICTAMEN_FINAL: HITO7_PLANTILLA_RESOLUCION_INSTITUCIONAL_GENERADA",
    ]
    out_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    expected = {
        "RESUMEN_EJECUTIVO",
        "CAMPOS_OFICIALES",
        "PROPUESTA_EQUIVALENCIAS",
        "CAMPOS_FORMULARIO",
        "CAMPOS_AUXILIARES",
        "MATRIZ_DECISION",
        "BASE_REFERENCIA",
        "PLAN_ACCION_HITO7",
        "CANDIDATO_AUX_SEPARADAS",
    }
    errors = []
    wb = load_workbook(out_xlsx, read_only=True, data_only=True)
    if set(wb.sheetnames) != expected:
        errors.append("Hojas Hito 7 no coinciden con lo esperado.")
    for path in [out_xlsx, out_md, out_official, out_form, out_json]:
        if not path.exists():
            errors.append(f"No existe salida requerida: {path}")
        try:
            path.relative_to(root / "cned")
        except ValueError:
            errors.append(f"Salida fuera de cned/: {path}")
    guard_ok_after, guard_after = run_guardrail(root)
    if not guard_ok_after:
        errors.append("Guardrail MU2026 fallo al final.")

    print("======================================================================")
    print("HITO 7 CNED - PLANTILLA RESOLUCION INSTITUCIONAL")
    print("======================================================================")
    print(f"Campos oficiales bloqueantes: {len(official)}")
    print(f"Equivalencia interna posible: {internal_possible}")
    print(f"Requieren fuente institucional: {require_source}")
    print(f"Campos formulario decision: {len(form)}")
    print(f"Campos auxiliares excluibles: {len(aux)}")
    print(f"Semaforo: {semaforo}")
    print(f"Excel: {out_xlsx}")
    print(f"Markdown: {out_md}")
    print(f"CSV oficiales: {out_official}")
    print(f"CSV formulario: {out_form}")
    print(f"JSON: {out_json}")
    print(f"Guardrail final: {guard_after}")
    if errors:
        print("AUTOAUDITORIA: ERROR")
        for error in errors:
            print(f"- {error}")
        return 1
    print("AUTOAUDITORIA: OK")
    print("DICTAMEN_FINAL: HITO7_PLANTILLA_RESOLUCION_INSTITUCIONAL_GENERADA")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
