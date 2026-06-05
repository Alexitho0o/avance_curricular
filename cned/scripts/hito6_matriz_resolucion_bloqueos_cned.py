"""Hito 6 CNED: matriz institucional de resolucion de bloqueos.

Consolida marcas y validaciones de Hito 5 en una matriz accionable para
resolucion institucional. No modifica bases ni genera carga final.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import unicodedata
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta pandas para ejecutar Hito 6 CNED.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta openpyxl para ejecutar Hito 6 CNED.") from exc


MARKS_REL = Path("cned/resultados/auditorias/CNED_HITO5_MARCAS_REVISION_20260605_135231.csv")
VALIDATIONS_REL = Path("cned/resultados/auditorias/CNED_HITO5_VALIDACIONES_REGISTROS_20260605_135231.csv")
BASE_HITO5_REL = Path("cned/resultados/archivos_subida/CNED_BASE_SANEADA_PRELIMINAR_HITO5_20260605_135231.xlsx")

OFFICIAL_BLOCKING = {
    "REGIMEN",
    "VIGENCIA_CARRERA",
    "ANIO_INICIO",
    "DURACION_REGIMEN",
    "NOMBRE_TITULO",
    "ACREDITACION",
    "REQUISITO_INGRESO",
    "SEMESTRES_RECONOCIDOS",
    "AREA_ACTUAL",
}

FORM_FIELD_HINTS = [
    "ARANCEL",
    "AREA",
    "SUB_AREA",
    "CARRERA_GENERICA",
    "DEPENDENCIA",
    "REQUISITO",
    "MALLA",
    "PERFIL",
    "MATRICULA",
    "COSTO",
    "CERTIFICADO",
    "DIPLOMA",
    "LICENCIA",
    "NOTAS",
    "PROMEDIO",
    "RECONOCIMIENTOS",
    "EXPERIENCIA",
    "MAIL",
    "FECHA_ADMISION",
    "INGRESO",
    "GRADO",
    "TITULO",
    "VACANTES",
    "CAMPUS",
    "HORARIO",
]


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def normalize(value: Any) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = text.strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text)
        if unicodedata.category(ch) != "Mn"
    )
    for char in [" ", ".", "/", "\\", "-", "(", ")", "[", "]", "\n", "\r", "\t", ":"]:
        text = text.replace(char, "_")
    text = re.sub(r"[^A-Z0-9_]+", "_", text)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def run_guardrail(root: Path) -> tuple[bool, str]:
    python_bin = root / ".venv" / "bin" / "python"
    cmd = [str(python_bin) if python_bin.exists() else "python", "cned/scripts/check_no_tocar_mu2026.py"]
    result = subprocess.run(cmd, cwd=root, text=True, capture_output=True)
    return result.returncode == 0, (result.stdout + result.stderr).strip()


def git_status(root: Path) -> str:
    result = subprocess.run(["git", "status", "--short", "--branch"], cwd=root, check=True, text=True, capture_output=True)
    return result.stdout.strip()


def max_severity(values: pd.Series) -> str:
    order = {"BAJA": 1, "MEDIA": 2, "ALTA": 3}
    seen = [str(v).upper() for v in values if str(v).upper() in order]
    if not seen:
        return "BAJA"
    return max(seen, key=lambda v: order[v])


def field_type(row: pd.Series) -> str:
    current = str(row.get("tipo_columna", ""))
    if current == "OFICIAL_PROBABLE":
        return "OFICIAL_PROBABLE"
    if current == "AUXILIAR_TECNICA":
        return "AUXILIAR_TECNICA"
    col = normalize(row.get("columna", ""))
    if any(hint in col for hint in FORM_FIELD_HINTS):
        return "FORMULARIO_CNED_DESCONOCIDO"
    return "FORMULARIO_CNED_DESCONOCIDO"


def resolution_category(campo: str, tipo: str, clasificacion_marca: str) -> str:
    norm = normalize(campo)
    mark = str(clasificacion_marca)
    if tipo == "AUXILIAR_TECNICA":
        return "EXCLUIR_DE_ARCHIVO_CARGA_SI_ES_AUXILIAR"
    if tipo == "OFICIAL_PROBABLE":
        if norm in {"ANIO_INICIO", "VIGENCIA_CARRERA", "DURACION_REGIMEN", "REGIMEN"}:
            return "VALIDAR_CON_REGISTRO_CURRICULAR"
        if norm in {"NOMBRE_TITULO", "ACREDITACION", "REQUISITO_INGRESO", "SEMESTRES_RECONOCIDOS", "AREA_ACTUAL"}:
            return "VALIDAR_CON_OFERTA_ACADEMICA"
        if "BLOQUEO" in mark:
            return "COMPLETAR_DESDE_FUENTE_OFICIAL"
        return "COMPLETAR_DESDE_FUENTE_OFICIAL"
    if any(token in norm for token in ["ARANCEL", "MATRICULA", "COSTO", "CERTIFICADO", "VALOR"]):
        return "VALIDAR_CON_FINANZAS_ARANCELES"
    if any(token in norm for token in ["MALLA", "PERFIL", "REQUISITO", "INGRESO", "LICENCIA", "NOTAS", "EXPERIENCIA"]):
        return "VALIDAR_CON_AREA_ACADEMICA"
    if any(token in norm for token in ["AREA", "SUB_AREA", "CARRERA_GENERICA", "DEPENDENCIA", "VACANTES"]):
        return "REQUIERE_DECISION_INSTITUCIONAL"
    if "OBSERVACION" in norm or norm == "OBSERVACIONES":
        return "MANTENER_SOLO_COMO_TRAZABILIDAD"
    return "NO_RESOLVER_AUTOMATICAMENTE"


def source_suggestion(campo: str, action: str, tipo: str) -> str:
    norm = normalize(campo)
    if tipo == "AUXILIAR_TECNICA":
        return "TRAZABILIDAD_INTERNA_CNED"
    if action == "VALIDAR_CON_FINANZAS_ARANCELES":
        return "FINANZAS_ARANCELES_INSTITUCIONAL"
    if action in {"VALIDAR_CON_AREA_ACADEMICA", "REQUIERE_DECISION_INSTITUCIONAL"}:
        return "AREA_ACADEMICA_O_DUENO_DEL_PROGRAMA"
    if action == "VALIDAR_CON_REGISTRO_CURRICULAR":
        return "REGISTRO_CURRICULAR_OFERTA_OFICIAL"
    if action == "VALIDAR_CON_OFERTA_ACADEMICA":
        return "OFERTA_ACADEMICA_INSTITUCIONAL"
    if norm in OFFICIAL_BLOCKING:
        return "FUENTE_OFICIAL_CNED_OFERTA_MATRIZ"
    return "INSTRUCTIVO_CNED_Y_DECISION_INSTITUCIONAL"


def blocks_upload(tipo: str, campo: str) -> str:
    norm = normalize(campo)
    if tipo == "AUXILIAR_TECNICA":
        return "NO"
    if tipo == "OFICIAL_PROBABLE" or norm in OFFICIAL_BLOCKING:
        return "SI"
    return "DEPENDE"


def can_exclude(tipo: str) -> str:
    if tipo == "AUXILIAR_TECNICA":
        return "SI"
    if tipo == "OFICIAL_PROBABLE":
        return "NO"
    return "DEPENDE"


def field_observation(tipo: str, campo: str, action: str) -> str:
    if tipo == "AUXILIAR_TECNICA":
        return "No debe ir en archivo de carga; conservar en trazabilidad."
    if tipo == "OFICIAL_PROBABLE":
        return "Campo estructural/oficial: no resolver sin fuente institucional."
    if action == "VALIDAR_CON_FINANZAS_ARANCELES":
        return "Requiere fuente monetaria explicita; no usar cero ni inferencias."
    return "Campo de formulario o carga potencial: validar contra instructivo CNED antes de excluir o completar."


def build_record_matrix(marks: pd.DataFrame, validations: pd.DataFrame) -> pd.DataFrame:
    name_map = {}
    if "NOMBRE_CARRERA" in validations.columns:
        for _, row in validations.iterrows():
            name_map[str(row.get("CODIGO_UNICO", ""))] = str(row.get("NOMBRE_CARRERA", ""))
    rows = []
    for _, row in marks.iterrows():
        tipo = field_type(row)
        campo = str(row.get("columna", ""))
        action = resolution_category(campo, tipo, str(row.get("clasificacion_marca", "")))
        codigo = str(row.get("CODIGO_UNICO", ""))
        rows.append({
            "fila_registro": row.get("fila_excel_aproximada", ""),
            "CODIGO_UNICO": codigo,
            "NOMBRE_CARRERA": name_map.get(codigo, ""),
            "campo": campo,
            "valor_actual": row.get("valor_completo", ""),
            "tipo_campo": tipo,
            "patron_detectado": row.get("patron_detectado", ""),
            "severidad": row.get("severidad_preliminar", ""),
            "motivo_bloqueo": row.get("clasificacion_marca", ""),
            "categoria_resolucion": action,
            "fuente_sugerida": source_suggestion(campo, action, tipo),
            "bloquea_subida": blocks_upload(tipo, campo),
            "se_puede_excluir": can_exclude(tipo),
            "observacion": field_observation(tipo, campo, action),
        })
    return pd.DataFrame(rows)


def build_field_summary(matrix: pd.DataFrame) -> pd.DataFrame:
    rows = []
    if matrix.empty:
        return pd.DataFrame()
    for campo, group in matrix.groupby("campo", dropna=False):
        tipo = group["tipo_campo"].mode().iloc[0]
        severity = max_severity(group["severidad"])
        action = group["categoria_resolucion"].mode().iloc[0]
        source = group["fuente_sugerida"].mode().iloc[0]
        rows.append({
            "campo": campo,
            "tipo_campo": tipo,
            "total_marcas": len(group),
            "total_registros_afectados": group["CODIGO_UNICO"].nunique(),
            "severidad_maxima": severity,
            "accion_recomendada": action,
            "fuente_sugerida": source,
            "bloquea_subida": blocks_upload(tipo, str(campo)),
            "se_puede_excluir": can_exclude(tipo),
            "observacion": field_observation(tipo, str(campo), action),
        })
    return pd.DataFrame(rows).sort_values(["bloquea_subida", "tipo_campo", "total_marcas"], ascending=[False, True, False])


def build_record_summary(matrix: pd.DataFrame, validations: pd.DataFrame) -> pd.DataFrame:
    rows = []
    by_code = matrix.groupby("CODIGO_UNICO") if not matrix.empty else []
    validation_map = {str(r.get("CODIGO_UNICO", "")): r for _, r in validations.iterrows()}
    for codigo, group in by_code:
        vrow = validation_map.get(str(codigo), {})
        rows.append({
            "CODIGO_UNICO": codigo,
            "NOMBRE_CARRERA": group["NOMBRE_CARRERA"].dropna().iloc[0] if not group["NOMBRE_CARRERA"].dropna().empty else "",
            "total_marcas": len(group),
            "campos_afectados": group["campo"].nunique(),
            "campos_oficiales_afectados": int((group["tipo_campo"] == "OFICIAL_PROBABLE").sum()),
            "campos_formulario_afectados": int((group["tipo_campo"] == "FORMULARIO_CNED_DESCONOCIDO").sum()),
            "campos_auxiliares_afectados": int((group["tipo_campo"] == "AUXILIAR_TECNICA").sum()),
            "estado_hito5": vrow.get("CNED_ESTADO_HITO5", ""),
            "alertas_hito5": vrow.get("CNED_ALERTAS_HITO5", ""),
            "puede_subir_hito5": vrow.get("CNED_PUEDE_SUBIR_PRELIMINAR_HITO5", ""),
        })
    return pd.DataFrame(rows)


def md_table(df: pd.DataFrame, columns: list[str], max_rows: int = 40) -> list[str]:
    if df.empty:
        return ["_Sin datos._"]
    cols = [c for c in columns if c in df.columns]
    subset = df.loc[:, cols].head(max_rows)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
    for _, row in subset.iterrows():
        lines.append("| " + " | ".join(str(row[c]).replace("|", "/") for c in cols) + " |")
    return lines


def write_csv(path: Path, df: pd.DataFrame) -> None:
    df.to_csv(path, index=False, quoting=csv.QUOTE_MINIMAL)


def safe(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy().astype(str)


def main() -> int:
    root = repo_root()
    guard_ok, guard_output = run_guardrail(root)
    if not guard_ok:
        print(guard_output)
        return 2
    marks_path = root / MARKS_REL
    validations_path = root / VALIDATIONS_REL
    base_path = root / BASE_HITO5_REL
    for path in [marks_path, validations_path, base_path]:
        if not path.exists():
            print(f"ERROR: falta insumo Hito 6: {path}")
            return 1

    marks = pd.read_csv(marks_path, dtype=str, keep_default_na=False)
    validations = pd.read_csv(validations_path, dtype=str, keep_default_na=False)
    base_ref = pd.read_excel(base_path, sheet_name="CANDIDATO_CARGA_LIMPIO", dtype=str, keep_default_na=False)
    matrix = build_record_matrix(marks, validations)
    field_summary = build_field_summary(matrix)
    record_summary = build_record_summary(matrix, validations)

    official_blockers = field_summary[field_summary["tipo_campo"] == "OFICIAL_PROBABLE"].copy()
    form_decision = field_summary[field_summary["tipo_campo"] == "FORMULARIO_CNED_DESCONOCIDO"].copy()
    aux_excludable = field_summary[field_summary["tipo_campo"] == "AUXILIAR_TECNICA"].copy()
    total_records_affected = matrix["CODIGO_UNICO"].nunique() if not matrix.empty else 0
    total_fields = matrix["campo"].nunique() if not matrix.empty else 0
    semaforo = "ROJO" if not official_blockers.empty else ("AMARILLO" if not form_decision.empty else "VERDE")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = root / "cned" / "resultados" / "archivos_subida" / f"CNED_MATRIZ_RESOLUCION_BLOQUEOS_HITO6_{timestamp}.xlsx"
    out_md = root / "cned" / "resultados" / "reportes" / f"CNED_HITO6_MATRIZ_RESOLUCION_BLOQUEOS_{timestamp}.md"
    out_field_csv = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO6_RESUMEN_POR_CAMPO_{timestamp}.csv"
    out_matrix_csv = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO6_MATRIZ_REGISTRO_CAMPO_{timestamp}.csv"
    out_json = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO6_RESUMEN_{timestamp}.json"
    for path in [out_xlsx, out_md, out_field_csv, out_matrix_csv, out_json]:
        path.parent.mkdir(parents=True, exist_ok=True)

    plan = pd.DataFrame([
        {"orden": 1, "tarea": "Completar campos oficiales bloqueantes", "detalle": "Resolver campos OFICIAL_PROBABLE con fuente institucional trazable.", "responsable_sugerido": "Registro Curricular / Oferta Academica"},
        {"orden": 2, "tarea": "Definir obligatoriedad de campos formulario", "detalle": "Confirmar con instructivo CNED que campos desconocidos son requeridos, opcionales o excluibles.", "responsable_sugerido": "Equipo CNED institucional"},
        {"orden": 3, "tarea": "Separar columnas auxiliares de trazabilidad", "detalle": "Excluir AUXILIAR_TECNICA del archivo de carga, conservandolas en auditoria.", "responsable_sugerido": "Equipo de datos"},
        {"orden": 4, "tarea": "Validar contra instructivo CNED", "detalle": "No generar carga final hasta confirmar estructura oficial.", "responsable_sugerido": "Responsable CNED"},
        {"orden": 5, "tarea": "Volver a ejecutar Hito 5", "detalle": "Revalidar despues de resolver bloqueos oficiales y decisiones de formulario.", "responsable_sugerido": "Equipo de datos"},
        {"orden": 6, "tarea": "Preparar candidato de subida", "detalle": "Solo cuando no queden bloqueos oficiales y las decisiones de formulario esten cerradas.", "responsable_sugerido": "Equipo CNED"},
    ])

    resumen = pd.DataFrame([
        {"campo": "fecha_ejecucion", "valor": datetime.now().isoformat(timespec="seconds")},
        {"campo": "registros_afectados", "valor": total_records_affected},
        {"campo": "campos_con_bloqueo", "valor": total_fields},
        {"campo": "campos_oficiales_bloqueantes", "valor": len(official_blockers)},
        {"campo": "campos_formulario_decision", "valor": len(form_decision)},
        {"campo": "campos_auxiliares_excluibles", "valor": len(aux_excludable)},
        {"campo": "total_marcas", "valor": len(matrix)},
        {"campo": "semaforo", "valor": semaforo},
        {"campo": "advertencia", "valor": "No es archivo final CNED; matriz para resolucion institucional."},
    ])

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        safe(resumen).to_excel(writer, sheet_name="RESUMEN_EJECUTIVO", index=False)
        safe(matrix).to_excel(writer, sheet_name="MATRIZ_RESOLUCION", index=False)
        safe(field_summary).to_excel(writer, sheet_name="RESUMEN_POR_CAMPO", index=False)
        safe(record_summary).to_excel(writer, sheet_name="RESUMEN_POR_REGISTRO", index=False)
        safe(official_blockers).to_excel(writer, sheet_name="CAMPOS_OFICIALES", index=False)
        safe(form_decision).to_excel(writer, sheet_name="CAMPOS_FORMULARIO", index=False)
        safe(aux_excludable).to_excel(writer, sheet_name="CAMPOS_AUXILIARES", index=False)
        safe(plan).to_excel(writer, sheet_name="PLAN_ACCION", index=False)
        safe(base_ref).to_excel(writer, sheet_name="BASE_REFERENCIA_HITO5", index=False)

    write_csv(out_field_csv, field_summary)
    write_csv(out_matrix_csv, matrix)

    payload = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "git_status_inicial": git_status(root),
        "insumos": {
            "marcas_hito5": str(MARKS_REL),
            "validaciones_hito5": str(VALIDATIONS_REL),
            "base_hito5": str(BASE_HITO5_REL),
        },
        "registros_afectados": int(total_records_affected),
        "campos_con_bloqueo": int(total_fields),
        "campos_oficiales_bloqueantes": official_blockers["campo"].tolist(),
        "campos_formulario_decision": form_decision["campo"].tolist(),
        "campos_auxiliares_excluibles": aux_excludable["campo"].tolist(),
        "total_marcas": int(len(matrix)),
        "semaforo": semaforo,
        "archivos_generados": {
            "excel": str(out_xlsx),
            "markdown": str(out_md),
            "resumen_por_campo_csv": str(out_field_csv),
            "matriz_registro_campo_csv": str(out_matrix_csv),
            "json": str(out_json),
        },
        "guardrail_mu2026": guard_output,
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    md_lines = [
        "# CNED Hito 6 - Matriz de resolucion de bloqueos",
        "",
        f"- Fecha: {datetime.now().isoformat(timespec='seconds')}",
        f"- Registros afectados: {total_records_affected}",
        f"- Campos con bloqueo/marca: {total_fields}",
        f"- Campos oficiales bloqueantes: {len(official_blockers)}",
        f"- Campos formulario con decision: {len(form_decision)}",
        f"- Campos auxiliares excluibles: {len(aux_excludable)}",
        f"- Semaforo final: **{semaforo}**",
        "",
        "## Por que no se puede subir aun",
        "",
        "Existen marcas en campos oficiales probables y campos de formulario sin decision institucional. No corresponde completar ni excluir automaticamente sin instructivo/fuente oficial.",
        "",
        "## Campos oficiales bloqueantes",
        "",
        *md_table(official_blockers, ["campo", "total_marcas", "total_registros_afectados", "accion_recomendada", "fuente_sugerida", "bloquea_subida"], 80),
        "",
        "## Campos de formulario que requieren decision",
        "",
        *md_table(form_decision, ["campo", "total_marcas", "total_registros_afectados", "accion_recomendada", "fuente_sugerida", "bloquea_subida", "se_puede_excluir"], 80),
        "",
        "## Campos auxiliares excluibles",
        "",
        *md_table(aux_excludable, ["campo", "total_marcas", "accion_recomendada", "se_puede_excluir"], 80),
        "",
        "## Plan de accion",
        "",
        *md_table(plan, ["orden", "tarea", "detalle", "responsable_sugerido"], 20),
        "",
        "## Criterio para volver a Hito 5",
        "",
        "Volver a ejecutar Hito 5 cuando los campos oficiales bloqueantes tengan valor validado y cuando los campos de formulario hayan sido clasificados como requeridos, opcionales o excluibles segun instructivo CNED.",
        "",
        "## Guardrail MU2026/PES_READY",
        "",
        "```text",
        guard_output,
        "```",
        "",
        "> Advertencia: no es archivo final CNED.",
        "",
        "## Dictamen",
        "",
        "DICTAMEN_FINAL: HITO6_MATRIZ_RESOLUCION_BLOQUEOS_GENERADA",
    ]
    out_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    expected = {
        "RESUMEN_EJECUTIVO",
        "MATRIZ_RESOLUCION",
        "RESUMEN_POR_CAMPO",
        "RESUMEN_POR_REGISTRO",
        "CAMPOS_OFICIALES",
        "CAMPOS_FORMULARIO",
        "CAMPOS_AUXILIARES",
        "PLAN_ACCION",
        "BASE_REFERENCIA_HITO5",
    }
    errors = []
    wb = load_workbook(out_xlsx, read_only=True, data_only=True)
    if set(wb.sheetnames) != expected:
        errors.append("Hojas Excel Hito 6 no coinciden con lo esperado.")
    if total_records_affected != validations["CODIGO_UNICO"].nunique():
        errors.append("Total registros afectados no coincide con validaciones Hito 5.")
    for path in [out_xlsx, out_md, out_field_csv, out_matrix_csv, out_json]:
        if not path.exists():
            errors.append(f"No existe salida requerida: {path}")
        try:
            path.relative_to(root / "cned")
        except ValueError:
            errors.append(f"Salida fuera de cned/: {path}")
    guard_ok_after, guard_after = run_guardrail(root)
    if not guard_ok_after:
        errors.append("Guardrail MU2026 fallo al final.")

    print("======================================================================")
    print("HITO 6 CNED - MATRIZ RESOLUCION BLOQUEOS")
    print("======================================================================")
    print(f"Registros afectados: {total_records_affected}")
    print(f"Campos con bloqueo: {total_fields}")
    print(f"Campos oficiales bloqueantes: {len(official_blockers)}")
    print(f"Campos formulario decision: {len(form_decision)}")
    print(f"Campos auxiliares excluibles: {len(aux_excludable)}")
    print(f"Semaforo: {semaforo}")
    print(f"Excel: {out_xlsx}")
    print(f"Markdown: {out_md}")
    print(f"CSV resumen campo: {out_field_csv}")
    print(f"CSV matriz: {out_matrix_csv}")
    print(f"JSON: {out_json}")
    print(f"Guardrail final: {guard_after}")
    if errors:
        print("AUTOAUDITORIA: ERROR")
        for error in errors:
            print(f"- {error}")
        return 1
    print("AUTOAUDITORIA: OK")
    print("DICTAMEN_FINAL: HITO6_MATRIZ_RESOLUCION_BLOQUEOS_GENERADA")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
