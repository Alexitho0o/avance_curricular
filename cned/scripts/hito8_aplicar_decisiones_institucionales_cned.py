"""Hito 8 CNED: aplicar decisiones institucionales trazables.

Lee la plantilla Hito 7 y aplica solo decisiones explicitas. Si la matriz
permanece pendiente, genera auditoria sin modificar valores.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta pandas para ejecutar Hito 8 CNED.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta openpyxl para ejecutar Hito 8 CNED.") from exc


HITO7_XLSX = Path("cned/resultados/archivos_subida/CNED_HITO7_PLANTILLA_RESOLUCION_INSTITUCIONAL_20260605_141142.xlsx")

REQUIRED_DECISION_COLUMNS = [
    "CAMPO",
    "TIPO_CAMPO",
    "PROPUESTA_ACCION",
    "FUENTE_REQUERIDA",
    "DECISION_INSTITUCIONAL",
    "VALOR_CORREGIDO_SI_APLICA",
    "OBSERVACION",
    "ESTADO_DECISION",
]

APPLICABLE = {
    "RESUELTO_CON_EQUIVALENCIA_INTERNA",
    "RESUELTO_CON_FUENTE_OFICIAL",
    "EXCLUIR_DE_CARGA",
    "MANTENER_SOLO_TRAZABILIDAD",
    "NO_APLICA",
}

NOT_APPLICABLE = {"", "PENDIENTE", "REQUIERE_COMITE"}

EQUIVALENCES = {
    "VIGENCIA_CARRERA": ["VIGENCIA"],
    "ANIO_INICIO": ["ANIO_INICIO", "Año de inicio de Actividades"],
    "DURACION_REGIMEN": ["Duración del programa en Semestres", "DURACION_REGIMEN"],
    "NOMBRE_TITULO": ["Título que otorga el programa", "NOMBRE_TITULO"],
    "REGIMEN": ["REGIMEN", "Régimen"],
    "Régimen": ["REGIMEN"],
    "ACREDITACION": ["ACREDITACION"],
    "AREA_ACTUAL": ["AREA_ACTUAL", "Área del Conocimiento"],
    "REQUISITO_INGRESO": ["REQUISITO_INGRESO", "TEXTO_REQUISITO_INGRESO"],
    "SEMESTRES_RECONOCIDOS": ["SEMESTRES_RECONOCIDOS"],
}

TECH_PATTERNS = [
    "REVISAR",
    "CONFLICTO",
    "DECISION",
    "PENDIENTE",
    "VALIDAR",
    "VERIFICAR",
    "DUDOSO",
    "SIN_FUENTE",
    "SIN_DECISION",
    "NO_DEFINIDO",
    "POR_DEFINIR",
]

CRITICAL_FIELDS = [
    "codigo_unico",
    "codigo_carrera",
    "nombre_carrera",
    "sede",
    "modalidad",
    "jornada",
    "version",
    "vigencia",
    "duracion",
]

KEYS = {
    "codigo_unico": ["CODIGO_UNICO", "COD_UNICO"],
    "codigo_sede": ["CODIGO_SEDE", "COD_SED", "COD_SEDE"],
    "codigo_carrera": ["CODIGO_CARRERA", "COD_CARRERA", "COD_CAR", "COD_CARRERA_DERIVADO"],
    "nombre_carrera": ["NOMBRE_CARRERA", "NOMBRE_CARRERA_ORIGINAL", "NOMBRE_CARRERA_NORMALIZADO", "NOMBRE"],
    "sede": ["CODIGO_SEDE", "COD_SED", "COD_SEDE", "NOMBRE_SEDE", "SEDE"],
    "modalidad": ["MODALIDAD", "MODALIDAD_ORIGINAL", "MODALIDAD_DEL_PROGRAMA"],
    "jornada": ["JORNADA", "CODIGO_JORNADA", "COD_JORNADA", "JOR"],
    "version": ["VERSION", "VERSION_CARRERA", "VERSION_DERIVADA"],
    "vigencia": ["VIGENCIA", "VIGENCIA_CARRERA"],
    "duracion": ["DURACION_TOTAL", "DURACION_ESTUDIOS", "DURACION_ESTUDIO", "DURACION_TITULACION", "DURACION_DEL_PROGRAMA_EN_SEMESTRES"],
}


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def normalize(value: Any) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = text.strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text)
        if unicodedata.category(ch) != "Mn"
    )
    for char in [" ", ".", "/", "\\", "-", "(", ")", "[", "]", "\n", "\r", "\t", ":"]:
        text = text.replace(char, "_")
    text = re.sub(r"[^A-Z0-9_]+", "_", text)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def value(value_: Any) -> str:
    if value_ is None or pd.isna(value_):
        return ""
    text = str(value_).strip()
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            return text
    return text


def contains_mark(value_: Any) -> bool:
    text = normalize(value_)
    return any(pattern in text for pattern in TECH_PATTERNS)


def run_guardrail(root: Path) -> tuple[bool, str]:
    py = root / ".venv" / "bin" / "python"
    cmd = [str(py) if py.exists() else "python", "cned/scripts/check_no_tocar_mu2026.py"]
    result = subprocess.run(cmd, cwd=root, text=True, capture_output=True)
    return result.returncode == 0, (result.stdout + result.stderr).strip()


def git_status(root: Path) -> str:
    result = subprocess.run(["git", "status", "--short", "--branch"], cwd=root, check=True, text=True, capture_output=True)
    return result.stdout.strip()


def colmap(df: pd.DataFrame) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for col in df.columns:
        mapping.setdefault(normalize(col), str(col))
    return mapping


def find_col(df: pd.DataFrame, names: list[str]) -> str | None:
    mapping = colmap(df)
    for name in names:
        col = mapping.get(normalize(name))
        if col:
            return col
    return None


def get_key(row: pd.Series, df: pd.DataFrame, key: str) -> str:
    col = find_col(df, KEYS.get(key, []))
    return value(row[col]) if col else ""


def parse_codigo_unico(code: str) -> dict[str, str] | None:
    match = re.fullmatch(r"I162S(?P<sede>\d+)C(?P<carrera>\d+)J(?P<jornada>\d+)V(?P<version>\d+)", code)
    return match.groupdict() if match else None


def is_positive_number(text: str) -> bool:
    try:
        return float(text) > 0
    except ValueError:
        return False


def load_base(workbook: Path) -> tuple[str, pd.DataFrame]:
    sheets = load_workbook(workbook, read_only=True, data_only=True).sheetnames
    for sheet in ["CANDIDATO_CON_AUXILIARES_SEPARADAS", "CANDIDATO_AUX_SEPARADAS", "BASE_REFERENCIA"]:
        if sheet in sheets:
            return sheet, pd.read_excel(workbook, sheet_name=sheet, dtype=str, keep_default_na=False)
    raise SystemExit("ERROR: no existe CANDIDATO_CON_AUXILIARES_SEPARADAS, CANDIDATO_AUX_SEPARADAS ni BASE_REFERENCIA.")


def classify_decisions(matrix: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows_app = []
    rows_no = []
    for _, row in matrix.iterrows():
        state = normalize(row.get("ESTADO_DECISION", ""))
        row_dict = row.to_dict()
        if state in APPLICABLE:
            rows_app.append(row_dict)
        else:
            if state not in NOT_APPLICABLE:
                row_dict["MOTIVO_NO_APLICA"] = f"ESTADO_NO_RECONOCIDO: {state}"
            elif not state:
                row_dict["MOTIVO_NO_APLICA"] = "ESTADO_DECISION_VACIO"
            else:
                row_dict["MOTIVO_NO_APLICA"] = state
            rows_no.append(row_dict)
    return pd.DataFrame(rows_app), pd.DataFrame(rows_no)


def apply_decisions(base: pd.DataFrame, applicable: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    corrected = base.copy()
    excluded_columns = []
    traces = []
    for _, decision in applicable.iterrows():
        field = str(decision.get("CAMPO", ""))
        state = normalize(decision.get("ESTADO_DECISION", ""))
        explicit = value(decision.get("VALOR_CORREGIDO_SI_APLICA", ""))
        source = str(decision.get("FUENTE_REQUERIDA", ""))
        observation = str(decision.get("OBSERVACION", ""))
        target = find_col(corrected, [field])

        if state == "EXCLUIR_DE_CARGA" or state == "MANTENER_SOLO_TRAZABILIDAD":
            if target:
                excluded_columns.append({
                    "CAMPO": field,
                    "COLUMNA_EXCLUIDA": target,
                    "ESTADO_DECISION": state,
                    "FUENTE_REQUERIDA": source,
                    "OBSERVACION": observation,
                })
            continue
        if state == "NO_APLICA":
            traces.append({
                "CAMPO_DESTINO": field,
                "CAMPO_FUENTE": "",
                "FILA": "",
                "VALOR_ANTERIOR": "",
                "VALOR_NUEVO": "",
                "ESTADO_DECISION": state,
                "FUENTE": source,
                "OBSERVACION": "Decision NO_APLICA registrada; no modifica valor.",
            })
            continue
        if not target:
            traces.append({
                "CAMPO_DESTINO": field,
                "CAMPO_FUENTE": "",
                "FILA": "",
                "VALOR_ANTERIOR": "",
                "VALOR_NUEVO": "",
                "ESTADO_DECISION": state,
                "FUENTE": source,
                "OBSERVACION": "No existe columna destino; no se aplica.",
            })
            continue
        if state == "RESUELTO_CON_FUENTE_OFICIAL":
            if not explicit:
                traces.append({
                    "CAMPO_DESTINO": field,
                    "CAMPO_FUENTE": "",
                    "FILA": "",
                    "VALOR_ANTERIOR": "",
                    "VALOR_NUEVO": "",
                    "ESTADO_DECISION": state,
                    "FUENTE": source,
                    "OBSERVACION": "Falta VALOR_CORREGIDO_SI_APLICA; no se aplica.",
                })
                continue
            for idx in corrected.index:
                old = value(corrected.at[idx, target])
                corrected.at[idx, target] = explicit
                traces.append({
                    "CAMPO_DESTINO": target,
                    "CAMPO_FUENTE": "VALOR_CORREGIDO_SI_APLICA",
                    "FILA": idx + 2,
                    "VALOR_ANTERIOR": old,
                    "VALOR_NUEVO": explicit,
                    "ESTADO_DECISION": state,
                    "FUENTE": source,
                    "OBSERVACION": observation,
                })
        elif state == "RESUELTO_CON_EQUIVALENCIA_INTERNA":
            source_col = None
            for candidate in EQUIVALENCES.get(field, []):
                col = find_col(corrected, [candidate])
                if col and col != target:
                    values = [value(v) for v in corrected[col].tolist()]
                    clean_values = [v for v in values if v and not contains_mark(v)]
                    if clean_values:
                        source_col = col
                        break
            if not source_col:
                traces.append({
                    "CAMPO_DESTINO": target,
                    "CAMPO_FUENTE": "",
                    "FILA": "",
                    "VALOR_ANTERIOR": "",
                    "VALOR_NUEVO": "",
                    "ESTADO_DECISION": state,
                    "FUENTE": source,
                    "OBSERVACION": "No existe equivalencia interna inequivoca; no se aplica.",
                })
                continue
            for idx in corrected.index:
                new = value(corrected.at[idx, source_col])
                old = value(corrected.at[idx, target])
                if new and not contains_mark(new):
                    corrected.at[idx, target] = new
                    traces.append({
                        "CAMPO_DESTINO": target,
                        "CAMPO_FUENTE": source_col,
                        "FILA": idx + 2,
                        "VALOR_ANTERIOR": old,
                        "VALOR_NUEVO": new,
                        "ESTADO_DECISION": state,
                        "FUENTE": source,
                        "OBSERVACION": observation,
                    })
    excluded = pd.DataFrame(excluded_columns)
    trace = pd.DataFrame(traces)
    return corrected, trace, excluded


def validate_records(candidate: pd.DataFrame, pending_official: list[str], undecided_form: list[str]) -> tuple[pd.DataFrame, pd.DataFrame]:
    code_col = find_col(candidate, KEYS["codigo_unico"])
    duplicate_codes = set()
    if code_col:
        codes = candidate[code_col].map(value)
        duplicate_codes = set(codes[(codes != "") & codes.duplicated(keep=False)].tolist())

    rows = []
    states = []
    alerts_all = []
    puede_all = []
    official_pending_existing = [field for field in pending_official if find_col(candidate, [field])]
    form_pending_existing = [field for field in undecided_form if find_col(candidate, [field])]

    for idx, row in candidate.iterrows():
        alerts = []
        critical = []
        code = get_key(row, candidate, "codigo_unico")
        if not code:
            alerts.append("CODIGO_UNICO_VACIO")
            critical.append("CODIGO_UNICO_VACIO")
        elif code in duplicate_codes:
            alerts.append("CODIGO_UNICO_DUPLICADO")
            critical.append("CODIGO_UNICO_DUPLICADO")
        parsed = parse_codigo_unico(code) if code else None
        if code and parsed is None:
            alerts.append("CODIGO_UNICO_MALFORMADO")
            critical.append("CODIGO_UNICO_MALFORMADO")
        for key in CRITICAL_FIELDS:
            if not get_key(row, candidate, key):
                label = f"{key.upper()}_VACIO"
                alerts.append(label)
                critical.append(label)
        dur = get_key(row, candidate, "duracion")
        if dur and not is_positive_number(dur):
            alerts.append("DURACION_INVALIDA")
            critical.append("DURACION_INVALIDA")
        if parsed:
            comparisons = {
                "sede": get_key(row, candidate, "codigo_sede"),
                "carrera": get_key(row, candidate, "codigo_carrera"),
                "jornada": get_key(row, candidate, "jornada"),
                "version": get_key(row, candidate, "version"),
            }
            for comp, separate in comparisons.items():
                if separate and parsed[comp] != separate:
                    label = f"CODIGO_UNICO_INCONSISTENTE_{comp.upper()}"
                    alerts.append(label)
                    critical.append(label)
        for field in official_pending_existing:
            col = find_col(candidate, [field])
            current = value(row[col]) if col else ""
            if contains_mark(current):
                alerts.append(f"CAMPO_OFICIAL_PENDIENTE_{normalize(field)}")
                critical.append(f"CAMPO_OFICIAL_PENDIENTE_{normalize(field)}")
        for field in form_pending_existing:
            col = find_col(candidate, [field])
            current = value(row[col]) if col else ""
            if contains_mark(current):
                alerts.append(f"CAMPO_FORMULARIO_PENDIENTE_{normalize(field)}")
        if critical:
            estado = "NO_SUBIR_BLOQUEO_CRITICO"
            puede = "NO"
        elif alerts:
            estado = "REQUIERE_REVISION"
            puede = "REVISAR"
        else:
            estado = "PUEDE_SUBIR_PRELIMINAR"
            puede = "SI"
        rows.append({
            "fila_excel_aproximada": idx + 2,
            "CODIGO_UNICO": code,
            "CNED_ESTADO_HITO8": estado,
            "CNED_ALERTAS_HITO8": " | ".join(sorted(set(alerts))),
            "CNED_PUEDE_SUBIR_PRELIMINAR_HITO8": puede,
        })
        states.append(estado)
        alerts_all.append(" | ".join(sorted(set(alerts))))
        puede_all.append(puede)
    out = candidate.copy()
    out["CNED_ESTADO_HITO8"] = states
    out["CNED_ALERTAS_HITO8"] = alerts_all
    out["CNED_PUEDE_SUBIR_PRELIMINAR_HITO8"] = puede_all
    return out, pd.DataFrame(rows)


def safe(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy().astype(str)


def write_csv(path: Path, df: pd.DataFrame) -> None:
    df.to_csv(path, index=False, quoting=csv.QUOTE_MINIMAL)


def md_table(df: pd.DataFrame, columns: list[str], max_rows: int = 50) -> list[str]:
    if df.empty:
        return ["_Sin datos._"]
    cols = [c for c in columns if c in df.columns]
    subset = df.loc[:, cols].head(max_rows)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
    for _, row in subset.iterrows():
        lines.append("| " + " | ".join(str(row[c]).replace("|", "/") for c in cols) + " |")
    return lines


def main() -> int:
    root = repo_root()
    guard_ok, guard_output = run_guardrail(root)
    if not guard_ok:
        print(guard_output)
        return 2
    hito7 = root / HITO7_XLSX
    if not hito7.exists():
        print(f"ERROR: no existe plantilla Hito 7: {hito7}")
        return 1
    wb = load_workbook(hito7, read_only=True, data_only=True)
    if "MATRIZ_DECISION" not in wb.sheetnames:
        print("ERROR: falta hoja MATRIZ_DECISION.")
        return 1
    matrix = pd.read_excel(hito7, sheet_name="MATRIZ_DECISION", dtype=str, keep_default_na=False)
    missing = [col for col in REQUIRED_DECISION_COLUMNS if col not in matrix.columns]
    if missing:
        print(f"ERROR: faltan columnas en MATRIZ_DECISION: {missing}")
        return 1
    base_sheet, base = load_base(hito7)
    applicable, not_applicable = classify_decisions(matrix)
    corrected, trace, excluded = apply_decisions(base, applicable)

    pending_official = matrix[
        (matrix["TIPO_CAMPO"].astype(str) == "OFICIAL_PROBABLE")
        & (~matrix["ESTADO_DECISION"].map(normalize).isin(APPLICABLE))
    ]["CAMPO"].astype(str).tolist()
    undecided_form = matrix[
        (matrix["TIPO_CAMPO"].astype(str) == "FORMULARIO_CNED_DESCONOCIDO")
        & (~matrix["ESTADO_DECISION"].map(normalize).isin(APPLICABLE))
    ]["CAMPO"].astype(str).tolist()

    excluded_cols = excluded["COLUMNA_EXCLUIDA"].tolist() if not excluded.empty and "COLUMNA_EXCLUIDA" in excluded.columns else []
    candidate_clean = corrected.drop(columns=excluded_cols, errors="ignore").copy()
    candidate_validated, validations = validate_records(candidate_clean, pending_official, undecided_form)
    si_df = candidate_validated[candidate_validated["CNED_PUEDE_SUBIR_PRELIMINAR_HITO8"] == "SI"].copy()
    revisar_df = candidate_validated[candidate_validated["CNED_PUEDE_SUBIR_PRELIMINAR_HITO8"] == "REVISAR"].copy()
    no_df = candidate_validated[candidate_validated["CNED_PUEDE_SUBIR_PRELIMINAR_HITO8"] == "NO"].copy()

    total = len(candidate_validated)
    si = len(si_df)
    revisar = len(revisar_df)
    no = len(no_df)
    semaforo = "VERDE" if si == total and total else ("ROJO" if no else "AMARILLO")
    recommendation = {
        "VERDE": "preparar carga candidata",
        "AMARILLO": "resolver revision",
        "ROJO": "no subir",
    }[semaforo]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = root / "cned" / "resultados" / "archivos_subida" / f"CNED_BASE_CORREGIDA_PRELIMINAR_HITO8_{timestamp}.xlsx"
    out_csv = root / "cned" / "resultados" / "archivos_subida" / f"CNED_CANDIDATO_PUEDE_SUBIR_HITO8_{timestamp}.csv"
    out_tsv = root / "cned" / "resultados" / "archivos_subida" / f"CNED_CANDIDATO_PUEDE_SUBIR_HITO8_{timestamp}.tsv"
    out_md = root / "cned" / "resultados" / "reportes" / f"CNED_HITO8_APLICACION_DECISIONES_INSTITUCIONALES_{timestamp}.md"
    out_trace = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO8_TRAZABILIDAD_CAMBIOS_{timestamp}.csv"
    out_valid = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO8_VALIDACIONES_REGISTROS_{timestamp}.csv"
    out_json = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO8_RESUMEN_{timestamp}.json"
    for path in [out_xlsx, out_csv, out_tsv, out_md, out_trace, out_valid, out_json]:
        path.parent.mkdir(parents=True, exist_ok=True)

    empty_trace = pd.DataFrame(columns=["CAMPO_DESTINO", "CAMPO_FUENTE", "FILA", "VALOR_ANTERIOR", "VALOR_NUEVO", "ESTADO_DECISION", "FUENTE", "OBSERVACION"])
    empty_excluded = pd.DataFrame(columns=["CAMPO", "COLUMNA_EXCLUIDA", "ESTADO_DECISION", "FUENTE_REQUERIDA", "OBSERVACION"])
    resumen = pd.DataFrame([
        {"campo": "fecha_ejecucion", "valor": datetime.now().isoformat(timespec="seconds")},
        {"campo": "plantilla_hito7", "valor": str(HITO7_XLSX)},
        {"campo": "hoja_base_usada", "valor": base_sheet},
        {"campo": "decisiones_aplicadas", "valor": len(applicable)},
        {"campo": "decisiones_no_aplicadas", "valor": len(not_applicable)},
        {"campo": "total_registros", "valor": total},
        {"campo": "total_si", "valor": si},
        {"campo": "total_revisar", "valor": revisar},
        {"campo": "total_no", "valor": no},
        {"campo": "semaforo", "valor": semaforo},
        {"campo": "recomendacion", "valor": recommendation},
        {"campo": "advertencia", "valor": "No es archivo final CNED."},
    ])

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        safe(resumen).to_excel(writer, sheet_name="RESUMEN_EJECUTIVO", index=False)
        safe(base).to_excel(writer, sheet_name="BASE_ORIGINAL_REFERENCIA", index=False)
        safe(corrected).to_excel(writer, sheet_name="BASE_CORREGIDA_TRAZABLE", index=False)
        safe(candidate_validated).to_excel(writer, sheet_name="CANDIDATO_CARGA_LIMPIO", index=False)
        safe(si_df).to_excel(writer, sheet_name="SOLO_PUEDE_SUBIR", index=False)
        safe(revisar_df).to_excel(writer, sheet_name="REQUIERE_REVISION", index=False)
        safe(no_df).to_excel(writer, sheet_name="NO_SUBIR", index=False)
        safe(trace if not trace.empty else empty_trace).to_excel(writer, sheet_name="TRAZABILIDAD_CAMBIOS", index=False)
        safe(applicable).to_excel(writer, sheet_name="DECISIONES_APLICADAS", index=False)
        safe(not_applicable).to_excel(writer, sheet_name="DECISIONES_NO_APLICADAS", index=False)
        safe(excluded if not excluded.empty else empty_excluded).to_excel(writer, sheet_name="COLUMNAS_EXCLUIDAS", index=False)
        safe(validations).to_excel(writer, sheet_name="VALIDACIONES_HITO8", index=False)

    write_csv(out_trace, trace if not trace.empty else empty_trace)
    write_csv(out_valid, validations)
    if si:
        write_csv(out_csv, si_df)
        si_df.to_csv(out_tsv, index=False, sep="\t", quoting=csv.QUOTE_MINIMAL)

    alert_counts = []
    for alerts in validations["CNED_ALERTAS_HITO8"].tolist():
        for alert in str(alerts).split(" | "):
            if alert:
                alert_counts.append(alert)
    alert_summary = pd.Series(alert_counts).value_counts().head(30).to_dict() if alert_counts else {}

    payload = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "git_status_inicial": git_status(root),
        "plantilla_hito7": str(HITO7_XLSX),
        "hoja_base_usada": base_sheet,
        "decisiones_aplicadas": int(len(applicable)),
        "decisiones_no_aplicadas": int(len(not_applicable)),
        "total_si": int(si),
        "total_revisar": int(revisar),
        "total_no": int(no),
        "semaforo": semaforo,
        "recomendacion": recommendation,
        "bloqueos_remanentes": alert_summary,
        "archivos_generados": {
            "excel": str(out_xlsx),
            "csv_si": str(out_csv) if si else None,
            "tsv_si": str(out_tsv) if si else None,
            "markdown": str(out_md),
            "trazabilidad_csv": str(out_trace),
            "validaciones_csv": str(out_valid),
            "json": str(out_json),
        },
        "guardrail_mu2026": guard_output,
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    md_lines = [
        "# CNED Hito 8 - Aplicacion de decisiones institucionales",
        "",
        f"- Fecha: {datetime.now().isoformat(timespec='seconds')}",
        f"- Plantilla Hito 7: `{HITO7_XLSX}`",
        f"- Hoja base usada: `{base_sheet}`",
        f"- Decisiones aplicadas: {len(applicable)}",
        f"- Decisiones no aplicadas: {len(not_applicable)}",
        f"- Total SI: {si}",
        f"- Total REVISAR: {revisar}",
        f"- Total NO: {no}",
        f"- Semaforo: **{semaforo}**",
        f"- Recomendacion: {recommendation}",
        "",
        "## Decisiones aplicadas",
        "",
        *md_table(applicable, ["CAMPO", "TIPO_CAMPO", "ESTADO_DECISION", "VALOR_CORREGIDO_SI_APLICA"], 80),
        "",
        "## Decisiones no aplicadas",
        "",
        *md_table(not_applicable, ["CAMPO", "TIPO_CAMPO", "ESTADO_DECISION", "MOTIVO_NO_APLICA"], 80),
        "",
        "## Bloqueos remanentes principales",
        "",
    ]
    if alert_summary:
        md_lines.extend(f"- {k}: {v}" for k, v in alert_summary.items())
    else:
        md_lines.append("- Sin bloqueos remanentes.")
    md_lines.extend([
        "",
        "## Archivos generados",
        "",
        f"- Excel Hito 8: `{out_xlsx.relative_to(root)}`",
        f"- Trazabilidad CSV: `{out_trace.relative_to(root)}`",
        f"- Validaciones CSV: `{out_valid.relative_to(root)}`",
        f"- JSON resumen: `{out_json.relative_to(root)}`",
        f"- CSV candidato SI: `{out_csv.relative_to(root) if si else 'NO_GENERADO_SIN_REGISTROS_SI'}`",
        f"- TSV candidato SI: `{out_tsv.relative_to(root) if si else 'NO_GENERADO_SIN_REGISTROS_SI'}`",
        "",
        "## Guardrail MU2026/PES_READY",
        "",
        "```text",
        guard_output,
        "```",
        "",
        "> Advertencia: no es archivo final CNED.",
        "",
        "## Dictamen",
        "",
        "DICTAMEN_FINAL: HITO8_DECISIONES_INSTITUCIONALES_PROCESADAS",
    ])
    out_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    expected = {
        "RESUMEN_EJECUTIVO",
        "BASE_ORIGINAL_REFERENCIA",
        "BASE_CORREGIDA_TRAZABLE",
        "CANDIDATO_CARGA_LIMPIO",
        "SOLO_PUEDE_SUBIR",
        "REQUIERE_REVISION",
        "NO_SUBIR",
        "TRAZABILIDAD_CAMBIOS",
        "DECISIONES_APLICADAS",
        "DECISIONES_NO_APLICADAS",
        "COLUMNAS_EXCLUIDAS",
        "VALIDACIONES_HITO8",
    }
    errors = []
    wb_out = load_workbook(out_xlsx, read_only=True, data_only=True)
    if set(wb_out.sheetnames) != expected:
        errors.append("Hojas Hito 8 no coinciden con lo esperado.")
    if si + revisar + no != total:
        errors.append("SI + REVISAR + NO no suma total.")
    for path in [out_xlsx, out_md, out_trace, out_valid, out_json]:
        if not path.exists():
            errors.append(f"No existe salida requerida: {path}")
        try:
            path.relative_to(root / "cned")
        except ValueError:
            errors.append(f"Salida fuera de cned/: {path}")
    guard_ok_after, guard_after = run_guardrail(root)
    if not guard_ok_after:
        errors.append("Guardrail MU2026 fallo al final.")

    print("======================================================================")
    print("HITO 8 CNED - APLICACION DECISIONES INSTITUCIONALES")
    print("======================================================================")
    print(f"Decisiones aplicadas: {len(applicable)}")
    print(f"Decisiones no aplicadas: {len(not_applicable)}")
    print(f"Total SI: {si}")
    print(f"Total REVISAR: {revisar}")
    print(f"Total NO: {no}")
    print(f"Semaforo: {semaforo}")
    print(f"Excel: {out_xlsx}")
    print(f"Markdown: {out_md}")
    print(f"CSV trazabilidad: {out_trace}")
    print(f"CSV validaciones: {out_valid}")
    print(f"JSON: {out_json}")
    print(f"CSV candidato SI: {out_csv if si else 'NO_GENERADO'}")
    print(f"TSV candidato SI: {out_tsv if si else 'NO_GENERADO'}")
    print(f"Guardrail final: {guard_after}")
    if errors:
        print("AUTOAUDITORIA: ERROR")
        for error in errors:
            print(f"- {error}")
        return 1
    print("AUTOAUDITORIA: OK")
    print("DICTAMEN_FINAL: HITO8_DECISIONES_INSTITUCIONALES_PROCESADAS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
