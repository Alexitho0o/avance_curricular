"""Hito 5 CNED: resolver marcas tecnicas de revision.

Traslada marcas tecnicas a trazabilidad, genera una base saneada preliminar y
revalida registros sin inventar datos ni declarar carga final.
"""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import re
import subprocess
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    import pandas as pd
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta pandas para ejecutar Hito 5 CNED.") from exc

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("ERROR: falta openpyxl para ejecutar Hito 5 CNED.") from exc


LOGGER = logging.getLogger("hito5_cned")

BASE_RELATIVE = Path("cned/resultados/archivos_subida/CNED_BASE_OPERATIVA_AUDITADA_PRE_SUBIDA_20260605_133904.xlsx")
BASE_SHEET = "BASE_AUDITADA"

MARK_PATTERNS = [
    "REVISAR",
    "REVISION",
    "REVISIÓN",
    "PENDIENTE",
    "VALIDAR",
    "VERIFICAR",
    "DUDA",
    "DUDOSO",
    "CONFLICTO",
    "OBS",
    "OBSERVACION",
    "OBSERVACIÓN",
    "MANUAL",
    "DECISION",
    "DECISIÓN",
    "AJUSTAR",
    "CORREGIR",
    "SIN DECISION",
    "SIN_DECISION",
    "NO DEFINIDO",
    "NO_DEFINIDO",
    "POR DEFINIR",
    "POR_DEFINIR",
    "SANEADO",
    "TECNICO",
    "TÉCNICO",
    "METODOLOGIA",
    "METODOLOGÍA",
]

OFFICIAL_PROBABLE = {
    "CODIGO_UNICO",
    "COD_UNICO",
    "CODIGO_IES",
    "CODIGO_IES_NUM",
    "COD_IES",
    "CODIGO_SEDE",
    "COD_SED",
    "COD_SEDE",
    "NOMBRE_SEDE",
    "CODIGO_CARRERA",
    "COD_CARRERA",
    "COD_CAR",
    "COD_CARRERA_DERIVADO",
    "NOMBRE_CARRERA",
    "NOMBRE_CARRERA_ORIGINAL",
    "NOMBRE_CARRERA_NORMALIZADO",
    "MODALIDAD",
    "MODALIDAD_ORIGINAL",
    "CODIGO_JORNADA",
    "COD_JORNADA",
    "JORNADA",
    "VERSION",
    "VERSION_CARRERA",
    "VERSION_DERIVADA",
    "TIPO_CARRERA",
    "CODIGO_TIPO_CARRERA",
    "CARACTERISTICAS_TIPO_PLAN",
    "TIPO_PLAN_CARRERA_ORIGINAL",
    "DURACION_ESTUDIOS",
    "DURACION_ESTUDIO",
    "DURACION_TITULACION",
    "DURACION_TOTAL",
    "REGIMEN",
    "DURACION_REGIMEN",
    "NOMBRE_TITULO",
    "CODIGO_NIVEL_GLOBAL",
    "CODIGO_NIVEL_CARRERA",
    "NIVEL_GLOBAL",
    "NIVEL_CARRERA",
    "ANIO_INICIO",
    "AÑO_INICIO",
    "ANO_INICIO",
    "ACREDITACION",
    "REQUISITO_INGRESO",
    "SEMESTRES_RECONOCIDOS",
    "AREA",
    "AREA_ACTUAL",
    "VIGENCIA",
    "VIGENCIA_CARRERA",
}

TECHNICAL_NAME_MARKERS = [
    "CNED_",
    "ALERTA",
    "ALERTAS",
    "AUDITORIA",
    "REVISION",
    "REVISIÓN",
    "OBSERVACION_AUDITORIA",
    "DECISION",
    "COMENTARIO",
    "MOTIVO",
    "TRAZABILIDAD",
    "MATCH",
    "SCORE",
    "FUENTE",
    "METODOLOGIA",
]

KEY_EQUIVALENCES = {
    "codigo_unico": ["CODIGO_UNICO", "COD_UNICO"],
    "codigo_sede": ["CODIGO_SEDE", "COD_SED", "COD_SEDE"],
    "codigo_carrera": ["CODIGO_CARRERA", "COD_CARRERA", "COD_CAR", "COD_CARRERA_DERIVADO"],
    "nombre_carrera": ["NOMBRE_CARRERA", "NOMBRE_CARRERA_ORIGINAL", "NOMBRE_CARRERA_NORMALIZADO", "NOMBRE"],
    "sede": ["CODIGO_SEDE", "COD_SED", "COD_SEDE", "NOMBRE_SEDE", "SEDE"],
    "modalidad": ["MODALIDAD", "MODALIDAD_ORIGINAL", "MODALIDAD_DEL_PROGRAMA"],
    "jornada": ["JORNADA", "CODIGO_JORNADA", "COD_JORNADA", "JOR"],
    "version": ["VERSION", "VERSION_CARRERA", "VERSION_DERIVADA"],
    "vigencia": ["VIGENCIA", "VIGENCIA_CARRERA"],
    "duracion": ["DURACION_TOTAL", "DURACION_ESTUDIOS", "DURACION_ESTUDIO", "DURACION_TITULACION", "DURACION_DEL_PROGRAMA_EN_SEMESTRES"],
}


def repo_root() -> Path:
    return Path(__file__).resolve().parents[2]


def normalize_text(value: Any) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = text.strip().upper()
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text)
        if unicodedata.category(ch) != "Mn"
    )
    return text


def normalize_header(value: Any) -> str:
    text = normalize_text(value)
    for char in [" ", ".", "/", "\\", "-", "(", ")", "[", "]", "\n", "\r", "\t", ":"]:
        text = text.replace(char, "_")
    text = re.sub(r"[^A-Z0-9_]+", "_", text)
    while "__" in text:
        text = text.replace("__", "_")
    return text.strip("_")


def normalize_value(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except ValueError:
            return text
    return text


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def run_guardrail(root: Path) -> tuple[bool, str]:
    python_bin = root / ".venv" / "bin" / "python"
    cmd = [str(python_bin) if python_bin.exists() else "python", "cned/scripts/check_no_tocar_mu2026.py"]
    result = subprocess.run(cmd, cwd=root, text=True, capture_output=True)
    return result.returncode == 0, (result.stdout + result.stderr).strip()


def git_status(root: Path) -> str:
    result = subprocess.run(["git", "status", "--short", "--branch"], cwd=root, text=True, capture_output=True, check=True)
    return result.stdout.strip()


def colmap(df: pd.DataFrame) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for col in df.columns:
        norm = normalize_header(col)
        if norm and norm not in mapping:
            mapping[norm] = col
    return mapping


def find_column(df: pd.DataFrame, key: str) -> str | None:
    mapping = colmap(df)
    for candidate in KEY_EQUIVALENCES.get(key, []):
        norm = normalize_header(candidate)
        if norm in mapping:
            return mapping[norm]
    return None


def classify_column(col: str) -> str:
    norm = normalize_header(col)
    if norm in OFFICIAL_PROBABLE:
        return "OFICIAL_PROBABLE"
    if norm.startswith("CNED_") or any(marker in norm for marker in TECHNICAL_NAME_MARKERS):
        return "AUXILIAR_TECNICA"
    return "DESCONOCIDA"


def detect_pattern(value: Any) -> str | None:
    text = normalize_text(value)
    if not text:
        return None
    for pattern in MARK_PATTERNS:
        normalized_pattern = normalize_text(pattern)
        if normalized_pattern in text:
            if normalized_pattern == "TECNICO" and not any(
                context in text
                for context in [
                    "MARCA_TECNICA",
                    "AUXILIAR_TECNICA",
                    "TECNICA_DE_REVISION",
                    "CAMPO_TECNICO",
                    "TECNICO_",
                    "_TECNICO",
                ]
            ):
                continue
            return pattern
    return None


def severity(column_type: str, col_norm: str) -> str:
    if column_type == "AUXILIAR_TECNICA":
        return "BAJA"
    if "OBS" in col_norm or "OBSERVACION" in col_norm:
        return "MEDIA"
    if column_type == "OFICIAL_PROBABLE":
        return "ALTA"
    return "MEDIA"


def mark_category(column_type: str, col_norm: str, value: str) -> str:
    value_norm = normalize_text(value)
    if column_type == "AUXILIAR_TECNICA":
        return "MARCA_TECNICA_REMOVIBLE"
    if "OBS" in col_norm or "OBSERVACION" in col_norm:
        return "OBSERVACION_INSTITUCIONAL"
    if column_type == "OFICIAL_PROBABLE":
        if "CONFLICTO" in value_norm or "DECISION" in value_norm:
            return "BLOQUEO_CRITICO"
        return "CAMPO_PENDIENTE_REAL"
    return "DUDA_METODOLOGICA"


def is_positive_number(value: str) -> bool:
    try:
        return float(value) > 0
    except ValueError:
        return False


def parse_codigo_unico(value: str) -> dict[str, str] | None:
    match = re.fullmatch(r"I162S(?P<sede>\d+)C(?P<carrera>\d+)J(?P<jornada>\d+)V(?P<version>\d+)", value)
    return match.groupdict() if match else None


def first_value(row: pd.Series, df: pd.DataFrame, key: str) -> str:
    col = find_column(df, key)
    return normalize_value(row[col]) if col is not None else ""


def classify_all_columns(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for col in df.columns:
        norm = normalize_header(col)
        ctype = classify_column(col)
        rows.append({
            "columna": str(col),
            "columna_normalizada": norm,
            "tipo_columna": ctype,
            "se_excluye_de_candidato_limpio": "SI" if ctype == "AUXILIAR_TECNICA" else "NO",
        })
    return pd.DataFrame(rows)


def detect_marks(df: pd.DataFrame, column_classes: dict[str, str]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for row_idx, row in df.iterrows():
        for col in df.columns:
            value = normalize_value(row[col])
            pattern = detect_pattern(value)
            if not pattern:
                continue
            col_norm = normalize_header(col)
            ctype = column_classes[str(col)]
            rows.append({
                "fila_excel_aproximada": row_idx + 2,
                "columna": str(col),
                "columna_normalizada": col_norm,
                "valor_completo": value,
                "patron_detectado": pattern,
                "tipo_columna": ctype,
                "severidad_preliminar": severity(ctype, col_norm),
                "clasificacion_marca": mark_category(ctype, col_norm, value),
                "CODIGO_UNICO": first_value(row, df, "codigo_unico"),
            })
    return pd.DataFrame(rows)


def build_clean_candidate(df: pd.DataFrame, column_df: pd.DataFrame) -> pd.DataFrame:
    aux_cols = set(column_df.loc[column_df["tipo_columna"] == "AUXILIAR_TECNICA", "columna"].astype(str))
    keep_cols = [col for col in df.columns if str(col) not in aux_cols]
    return df.loc[:, keep_cols].copy()


def revalidate_candidate(candidate: pd.DataFrame, marks_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    official_mark_rows = marks_df[marks_df["tipo_columna"] == "OFICIAL_PROBABLE"] if not marks_df.empty else pd.DataFrame()
    unknown_mark_rows = marks_df[marks_df["tipo_columna"] == "DESCONOCIDA"] if not marks_df.empty else pd.DataFrame()
    official_by_code: defaultdict[str, list[str]] = defaultdict(list)
    unknown_by_code: defaultdict[str, list[str]] = defaultdict(list)
    if not official_mark_rows.empty:
        for _, row in official_mark_rows.iterrows():
            official_by_code[str(row["CODIGO_UNICO"])].append(f"{row['columna']}={row['valor_completo']}")
    if not unknown_mark_rows.empty:
        for _, row in unknown_mark_rows.iterrows():
            unknown_by_code[str(row["CODIGO_UNICO"])].append(f"{row['columna']}={row['valor_completo']}")

    code_col = find_column(candidate, "codigo_unico")
    duplicate_codes = set()
    if code_col:
        codes = candidate[code_col].map(normalize_value)
        duplicate_codes = set(codes[(codes != "") & codes.duplicated(keep=False)].tolist())

    validation_rows = []
    status_values = []
    alert_values = []
    puede_values = []
    for idx, row in candidate.iterrows():
        alerts: list[str] = []
        critical: list[str] = []
        codigo = first_value(row, candidate, "codigo_unico")
        if not codigo:
            alerts.append("CODIGO_UNICO_VACIO")
            critical.append("CODIGO_UNICO_VACIO")
        elif codigo in duplicate_codes:
            alerts.append("CODIGO_UNICO_DUPLICADO")
            critical.append("CODIGO_UNICO_DUPLICADO")
        parsed = parse_codigo_unico(codigo) if codigo else None
        if codigo and parsed is None:
            alerts.append("CODIGO_UNICO_MALFORMADO")
            critical.append("CODIGO_UNICO_MALFORMADO")

        for key, label in [
            ("codigo_carrera", "CODIGO_CARRERA_VACIO"),
            ("nombre_carrera", "NOMBRE_CARRERA_VACIO"),
            ("sede", "SEDE_VACIA"),
            ("modalidad", "MODALIDAD_VACIA"),
            ("jornada", "JORNADA_VACIA"),
            ("version", "VERSION_VACIA"),
            ("vigencia", "VIGENCIA_VACIA"),
            ("duracion", "DURACION_VACIA"),
        ]:
            if not first_value(row, candidate, key):
                alerts.append(label)
                critical.append(label)

        dur = first_value(row, candidate, "duracion")
        if dur and not is_positive_number(dur):
            alerts.append("DURACION_INVALIDA")
            critical.append("DURACION_INVALIDA")

        if parsed:
            comparisons = {
                "sede": first_value(row, candidate, "codigo_sede"),
                "carrera": first_value(row, candidate, "codigo_carrera"),
                "jornada": first_value(row, candidate, "jornada"),
                "version": first_value(row, candidate, "version"),
            }
            for component, separate_value in comparisons.items():
                if separate_value and parsed[component] != separate_value:
                    issue = f"CODIGO_UNICO_INCONSISTENTE_{component.upper()}"
                    alerts.append(issue)
                    critical.append(issue)

        if codigo in official_by_code:
            alerts.append("MARCAS_TECNICAS_EN_CAMPOS_OFICIALES")
            critical.append("MARCAS_TECNICAS_EN_CAMPOS_OFICIALES")
        if codigo in unknown_by_code:
            alerts.append("MARCAS_TECNICAS_EN_COLUMNAS_DESCONOCIDAS")

        if critical:
            puede = "NO"
            estado = "NO_SUBIR_BLOQUEO_CRITICO"
        elif alerts:
            puede = "REVISAR"
            estado = "REQUIERE_REVISION"
        else:
            puede = "SI"
            estado = "PUEDE_SUBIR_PRELIMINAR"
        status_values.append(estado)
        alert_values.append(" | ".join(sorted(set(alerts))))
        puede_values.append(puede)
        validation_rows.append({
            "fila_excel_aproximada": idx + 2,
            "CODIGO_UNICO": codigo,
            "CNED_ESTADO_HITO5": estado,
            "CNED_ALERTAS_HITO5": " | ".join(sorted(set(alerts))),
            "CNED_PUEDE_SUBIR_PRELIMINAR_HITO5": puede,
            "DETALLE_MARCAS_OFICIALES": " || ".join(official_by_code.get(codigo, [])),
            "DETALLE_MARCAS_DESCONOCIDAS": " || ".join(unknown_by_code.get(codigo, [])),
        })

    candidate = candidate.copy()
    candidate["CNED_ESTADO_HITO5"] = status_values
    candidate["CNED_ALERTAS_HITO5"] = alert_values
    candidate["CNED_PUEDE_SUBIR_PRELIMINAR_HITO5"] = puede_values
    return candidate, pd.DataFrame(validation_rows)


def profile_clean_columns(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    total = max(len(df), 1)
    for col in df.columns:
        values = df[col].map(normalize_value)
        non_empty = int((values != "").sum())
        rows.append({
            "columna": str(col),
            "columna_normalizada": normalize_header(col),
            "valores_no_vacios": non_empty,
            "valores_vacios": total - non_empty,
            "porcentaje_completitud": round(non_empty / total * 100, 2),
            "valores_unicos": len([v for v in values.unique().tolist() if v]),
            "tipo_columna": classify_column(col),
        })
    return pd.DataFrame(rows)


def write_csv(path: Path, df: pd.DataFrame, sep: str = ",") -> None:
    df.to_csv(path, index=False, sep=sep, quoting=csv.QUOTE_MINIMAL)


def safe(df: pd.DataFrame) -> pd.DataFrame:
    return df.copy().astype(str)


def md_table(df: pd.DataFrame, columns: list[str], max_rows: int = 30) -> list[str]:
    if df.empty:
        return ["_Sin datos._"]
    cols = [c for c in columns if c in df.columns]
    subset = df.loc[:, cols].head(max_rows)
    lines = ["| " + " | ".join(cols) + " |", "| " + " | ".join(["---"] * len(cols)) + " |"]
    for _, row in subset.iterrows():
        lines.append("| " + " | ".join(str(row[c]).replace("|", "/") for c in cols) + " |")
    return lines


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    root = repo_root()
    guard_ok, guard_output = run_guardrail(root)
    if not guard_ok:
        print(guard_output)
        return 2

    base_path = root / BASE_RELATIVE
    if not base_path.exists():
        LOGGER.error("No existe Excel auditado Hito 4: %s", base_path)
        return 1
    base_hash_before = sha256(base_path)

    try:
        wb = load_workbook(base_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        LOGGER.error("No se pudo abrir Excel auditado Hito 4: %s", exc)
        return 1
    if BASE_SHEET not in wb.sheetnames:
        LOGGER.error("No existe hoja requerida %s en %s", BASE_SHEET, base_path)
        return 1

    df = pd.read_excel(base_path, sheet_name=BASE_SHEET, dtype=str, keep_default_na=False)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx = root / "cned" / "resultados" / "archivos_subida" / f"CNED_BASE_SANEADA_PRELIMINAR_HITO5_{timestamp}.xlsx"
    out_csv = root / "cned" / "resultados" / "archivos_subida" / f"CNED_CANDIDATO_PUEDE_SUBIR_HITO5_{timestamp}.csv"
    out_tsv = root / "cned" / "resultados" / "archivos_subida" / f"CNED_CANDIDATO_PUEDE_SUBIR_HITO5_{timestamp}.tsv"
    out_md = root / "cned" / "resultados" / "reportes" / f"CNED_HITO5_RESOLUCION_MARCAS_REVISION_{timestamp}.md"
    out_marks = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO5_MARCAS_REVISION_{timestamp}.csv"
    out_validations = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO5_VALIDACIONES_REGISTROS_{timestamp}.csv"
    out_json = root / "cned" / "resultados" / "auditorias" / f"CNED_HITO5_RESUMEN_{timestamp}.json"
    for path in [out_xlsx, out_csv, out_tsv, out_md, out_marks, out_validations, out_json]:
        path.parent.mkdir(parents=True, exist_ok=True)

    columns_df = classify_all_columns(df)
    column_classes = dict(zip(columns_df["columna"], columns_df["tipo_columna"]))
    marks_df = detect_marks(df, column_classes)
    candidate_raw = build_clean_candidate(df, columns_df)
    candidate, validations_df = revalidate_candidate(candidate_raw, marks_df)
    clean_dict_df = profile_clean_columns(candidate)

    puede_df = candidate[candidate["CNED_PUEDE_SUBIR_PRELIMINAR_HITO5"] == "SI"].copy()
    revisar_df = candidate[candidate["CNED_PUEDE_SUBIR_PRELIMINAR_HITO5"] == "REVISAR"].copy()
    no_df = candidate[candidate["CNED_PUEDE_SUBIR_PRELIMINAR_HITO5"] == "NO"].copy()

    total = len(candidate)
    si = len(puede_df)
    revisar = len(revisar_df)
    no = len(no_df)
    marks_total = len(marks_df)
    official_marks = int((marks_df["tipo_columna"] == "OFICIAL_PROBABLE").sum()) if not marks_df.empty else 0
    aux_marks = int((marks_df["tipo_columna"] == "AUXILIAR_TECNICA").sum()) if not marks_df.empty else 0
    unknown_marks = int((marks_df["tipo_columna"] == "DESCONOCIDA").sum()) if not marks_df.empty else 0
    critical_marks = int((marks_df["clasificacion_marca"] == "BLOQUEO_CRITICO").sum()) if not marks_df.empty else 0

    if no > 0 or official_marks > 0:
        semaforo = "ROJO"
    elif total and si / total > 0.95:
        semaforo = "VERDE"
    else:
        semaforo = "AMARILLO"

    resumen_df = pd.DataFrame([
        {"campo": "fecha_ejecucion", "valor": datetime.now().isoformat(timespec="seconds")},
        {"campo": "base_hito4", "valor": str(BASE_RELATIVE)},
        {"campo": "sha256_base_antes", "valor": base_hash_before},
        {"campo": "total_registros", "valor": total},
        {"campo": "total_marcas_detectadas", "valor": marks_total},
        {"campo": "marcas_columnas_oficiales", "valor": official_marks},
        {"campo": "marcas_columnas_auxiliares", "valor": aux_marks},
        {"campo": "marcas_columnas_desconocidas", "valor": unknown_marks},
        {"campo": "bloqueos_criticos_por_marca", "valor": critical_marks},
        {"campo": "total_si", "valor": si},
        {"campo": "total_revisar", "valor": revisar},
        {"campo": "total_no", "valor": no},
        {"campo": "semaforo", "valor": semaforo},
        {"campo": "advertencia", "valor": "No es archivo final sin instructivo oficial confirmado."},
    ])

    empty_marks = pd.DataFrame(columns=[
        "fila_excel_aproximada", "columna", "columna_normalizada", "valor_completo",
        "patron_detectado", "tipo_columna", "severidad_preliminar",
        "clasificacion_marca", "CODIGO_UNICO",
    ])

    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        safe(resumen_df).to_excel(writer, sheet_name="RESUMEN_EJECUTIVO", index=False)
        safe(df).to_excel(writer, sheet_name="BASE_AUDITADA_ORIGINAL", index=False)
        safe(candidate).to_excel(writer, sheet_name="CANDIDATO_CARGA_LIMPIO", index=False)
        safe(puede_df).to_excel(writer, sheet_name="SOLO_PUEDE_SUBIR", index=False)
        safe(revisar_df).to_excel(writer, sheet_name="REQUIERE_REVISION", index=False)
        safe(no_df).to_excel(writer, sheet_name="NO_SUBIR", index=False)
        safe(marks_df if not marks_df.empty else empty_marks).to_excel(writer, sheet_name="TRAZABILIDAD_MARCAS", index=False)
        safe(columns_df).to_excel(writer, sheet_name="COLUMNAS_CLASIFICADAS", index=False)
        safe(validations_df).to_excel(writer, sheet_name="VALIDACIONES_HITO5", index=False)
        safe(clean_dict_df).to_excel(writer, sheet_name="DICCIONARIO_CANDIDATO", index=False)

    write_csv(out_marks, marks_df if not marks_df.empty else empty_marks)
    write_csv(out_validations, validations_df)
    if si:
        write_csv(out_csv, puede_df)
        write_csv(out_tsv, puede_df, sep="\t")

    marks_by_column = (
        marks_df.groupby(["columna", "tipo_columna", "clasificacion_marca"]).size().reset_index(name="marcas")
        if not marks_df.empty else pd.DataFrame(columns=["columna", "tipo_columna", "clasificacion_marca", "marcas"])
    )
    validation_alert_counter = Counter()
    for value in validations_df["CNED_ALERTAS_HITO5"].tolist():
        for alert in str(value).split(" | "):
            if alert:
                validation_alert_counter[alert] += 1

    payload = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "git_status_inicial": git_status(root),
        "base_hito4": str(BASE_RELATIVE),
        "sha256_base_antes": base_hash_before,
        "sha256_base_despues": sha256(base_path),
        "total_registros": total,
        "total_marcas_detectadas": marks_total,
        "marcas_columnas_oficiales": official_marks,
        "marcas_columnas_auxiliares": aux_marks,
        "marcas_columnas_desconocidas": unknown_marks,
        "total_si": si,
        "total_revisar": revisar,
        "total_no": no,
        "semaforo": semaforo,
        "columnas_con_marcas": marks_by_column.to_dict(orient="records"),
        "alertas_validacion": validation_alert_counter.most_common(20),
        "archivos_generados": {
            "excel": str(out_xlsx),
            "csv_si": str(out_csv) if si else None,
            "tsv_si": str(out_tsv) if si else None,
            "markdown": str(out_md),
            "marcas_csv": str(out_marks),
            "validaciones_csv": str(out_validations),
            "json": str(out_json),
        },
        "guardrail_mu2026": guard_output,
        "advertencia": "No es archivo final sin instructivo oficial confirmado.",
    }
    out_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    md_lines = [
        "# CNED Hito 5 - Resolucion de marcas tecnicas de revision",
        "",
        f"- Fecha: {datetime.now().isoformat(timespec='seconds')}",
        f"- Base usada: `{BASE_RELATIVE}`",
        f"- Total registros originales: {total}",
        f"- Total marcas detectadas: {marks_total}",
        f"- Marcas en columnas oficiales: {official_marks}",
        f"- Marcas en columnas auxiliares: {aux_marks}",
        f"- Marcas en columnas desconocidas: {unknown_marks}",
        f"- Total SI: {si}",
        f"- Total REVISAR: {revisar}",
        f"- Total NO: {no}",
        f"- Semaforo: **{semaforo}**",
        "",
        "## Estado Git inicial",
        "",
        "```text",
        git_status(root),
        "```",
        "",
        "## Columnas con marcas",
        "",
        *md_table(marks_by_column, ["columna", "tipo_columna", "clasificacion_marca", "marcas"], 80),
        "",
        "## Motivos principales Hito 5",
        "",
    ]
    if validation_alert_counter:
        md_lines.extend(f"- {alert}: {count}" for alert, count in validation_alert_counter.most_common(20))
    else:
        md_lines.append("- Sin alertas.")
    md_lines.extend([
        "",
        "## Cambio respecto a Hito 4",
        "",
        "Hito 4 marcaba los 71 registros en revision porque detectaba marcas tecnicas en cualquier columna. Hito 5 separa esas marcas por tipo de columna: auxiliares, oficiales probables y desconocidas. Las marcas auxiliares se trasladan a trazabilidad y no bloquean por si solas; las marcas oficiales bloquean; las desconocidas mantienen revision.",
        "",
        "## Archivos generados",
        "",
        f"- Excel Hito 5: `{out_xlsx.relative_to(root)}`",
        f"- Auditoria marcas: `{out_marks.relative_to(root)}`",
        f"- Auditoria validaciones: `{out_validations.relative_to(root)}`",
        f"- JSON resumen: `{out_json.relative_to(root)}`",
        f"- CSV candidato SI: `{out_csv.relative_to(root) if si else 'NO_GENERADO_SIN_REGISTROS_SI'}`",
        f"- TSV candidato SI: `{out_tsv.relative_to(root) if si else 'NO_GENERADO_SIN_REGISTROS_SI'}`",
        "",
        "## Guardrail MU2026/PES_READY",
        "",
        "```text",
        guard_output,
        "```",
        "",
        "> Advertencia: no es archivo final sin instructivo oficial confirmado.",
        "",
        "## Dictamen",
        "",
        "DICTAMEN_FINAL: HITO5_MARCAS_REVISION_RESUELTAS_CON_BASE_SANEADA_PRELIMINAR",
    ])
    out_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    expected = {
        "RESUMEN_EJECUTIVO",
        "BASE_AUDITADA_ORIGINAL",
        "CANDIDATO_CARGA_LIMPIO",
        "SOLO_PUEDE_SUBIR",
        "REQUIERE_REVISION",
        "NO_SUBIR",
        "TRAZABILIDAD_MARCAS",
        "COLUMNAS_CLASIFICADAS",
        "VALIDACIONES_HITO5",
        "DICCIONARIO_CANDIDATO",
    }
    self_errors = []
    if not out_xlsx.exists():
        self_errors.append("No existe Excel Hito 5.")
    else:
        out_wb = load_workbook(out_xlsx, read_only=True, data_only=True)
        if set(out_wb.sheetnames) != expected:
            self_errors.append("Hojas Excel Hito 5 no coinciden con lo esperado.")
    if si + revisar + no != total:
        self_errors.append("SI + REVISAR + NO no suma total.")
    if sha256(base_path) != base_hash_before:
        self_errors.append("La base Hito 4 cambio durante la ejecucion.")
    for path in [out_xlsx, out_md, out_marks, out_validations, out_json]:
        if not path.exists():
            self_errors.append(f"No existe salida requerida: {path}")
        try:
            path.relative_to(root / "cned")
        except ValueError:
            self_errors.append(f"Salida escrita fuera de cned/: {path}")
    guard_ok_after, guard_after = run_guardrail(root)
    if not guard_ok_after:
        self_errors.append("Guardrail MU2026 fallo al final.")

    print("======================================================================")
    print("HITO 5 CNED - RESOLUCION MARCAS REVISION")
    print("======================================================================")
    print(f"Total registros: {total}")
    print(f"Total marcas detectadas: {marks_total}")
    print(f"Marcas oficiales: {official_marks}")
    print(f"Marcas auxiliares: {aux_marks}")
    print(f"Marcas desconocidas: {unknown_marks}")
    print(f"Total SI: {si}")
    print(f"Total REVISAR: {revisar}")
    print(f"Total NO: {no}")
    print(f"Semaforo: {semaforo}")
    print(f"Excel Hito 5: {out_xlsx}")
    print(f"Markdown: {out_md}")
    print(f"CSV marcas: {out_marks}")
    print(f"CSV validaciones: {out_validations}")
    print(f"JSON resumen: {out_json}")
    print(f"CSV candidato SI: {out_csv if si else 'NO_GENERADO'}")
    print(f"TSV candidato SI: {out_tsv if si else 'NO_GENERADO'}")
    print(f"Guardrail final: {guard_after}")
    if self_errors:
        print("AUTOAUDITORIA: ERROR")
        for error in self_errors:
            print(f"- {error}")
        return 1
    print("AUTOAUDITORIA: OK")
    print("DICTAMEN_FINAL: HITO5_MARCAS_REVISION_RESUELTAS_CON_BASE_SANEADA_PRELIMINAR")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
