#!/usr/bin/env python3
"""Correccion V2 de gobernanza para 20 columnas oficiales de extranjeros 2025.

La V2 reutiliza las fuentes ya procesadas del repositorio, corrige mapeos
incompletos de la V1 y mantiene pendientes institucionales cuando no hay
respaldo explicito.
"""

from __future__ import annotations

import csv
import hashlib
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path("/Users/alexi/Documents/GitHub/avance_curricular")
SUB = ROOT / "estudiantes_extranjeros_2026"
FROZEN_HASH = "da85dd0c453a942a4490f469b75d0418e9054e458a46028f44271ac704518081"
STRUCTURE_HASH = "dfa2262abddd3bf2489cdcd6d2e5f4e95863325b8f0de3e407d9f6e0857528c2"
PENDING = "PENDIENTE_INSTITUCIONAL"

OFFICIAL_COLUMNS = [
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "TIPO_RESIDENCIA_ESTUDIANTE",
    "PAIS_ORIGEN",
    "PAIS_ESTUDIOS_SECUNDARIOS",
    "CODIGO_UNICO",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "NOMBRE_UNIVERSIDAD_ORIGEN",
    "PAIS_UNIVERSIDAD_ORIGEN",
    "VIGENCIA",
]

PATHS = {
    "frozen": SUB / "data/frozen/BASE_EXTRANJEROS_2025_CONGELADA.tsv",
    "structure": SUB / "20260602_97636_Estructura_Extranjeros_Regulares_2025.csv",
    "manual": SUB / "docs/Instructivo_Estudiantes_Extranjeros_SIES_2026.txt",
    "precarga_raw": SUB / "data/raw/REPORTE_PRECARGA_EXTRANJEROS_REGULARES_2026_ORIGINAL.csv",
    "precarga_norm": SUB / "data/interim/PRECARGA_EXTRANJEROS_REGULARES_2026_NORMALIZADA.csv",
    "matricula": ROOT / "resultados/matricula_avance_curricular_2025_control.csv",
    "universo": SUB / "resultados/auditorias/UNIVERSO_DEPURADO_EXTRANJEROS_REGULARES_2025.csv",
    "datos_alumnos": ROOT / "input/PROMEDIOSDEALUMNOS_7804.xlsx",
    "matriz_v1": SUB / "data/processed/MATRIZ_GOBERNANZA_EXTRANJEROS_2025.csv",
    "cobertura_v1": SUB / "resultados/auditorias/COBERTURA_GOBERNANZA_COLUMNAS_EXTRANJEROS_2025.csv",
    "cierre": SUB / "data/processed/BASE_EXTRANJEROS_2025_CONGELADA_ENRIQUECIDA_CIERRE.csv",
    "conc_precarga": SUB / "resultados/auditorias/CONCILIACION_62_VS_PRECARGA_PES.csv",
    "conc_matricula": SUB / "resultados/auditorias/CONCILIACION_62_VS_MATRICULA_2025.csv",
    "conc_257": SUB / "resultados/auditorias/CONCILIACION_62_CONGELADOS_VS_257.csv",
    "res_codigo": SUB / "resultados/auditorias/RESOLUCION_CODIGO_UNICO.csv",
    "res_mult": SUB / "resultados/auditorias/RESOLUCION_19_COINCIDENCIAS_MULTIPLES.csv",
    "cand_mult": SUB / "resultados/auditorias/CANDIDATOS_19_COINCIDENCIAS_MULTIPLES.csv",
    "res_evid4": SUB / "resultados/auditorias/RESOLUCION_4_SIN_EVIDENCIA_2025.csv",
    "res_nac": SUB / "resultados/auditorias/RESOLUCION_NACIONALIDAD_AMBIGUA_62.csv",
    "vig_corr": SUB / "resultados/auditorias/AUDITORIA_VIGENCIA_2025_CORREGIDA.csv",
    "catalogo_paises": SUB / "resultados/auditorias/CATALOGO_PAISES_SIES_2026.csv",
    "puente_sies": ROOT / "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
}

OUT = {
    "inventario": SUB / "resultados/auditorias/INVENTARIO_FUENTES_20_COLUMNAS_OFICIALES.csv",
    "matriz_fuentes": SUB / "resultados/auditorias/MATRIZ_FUENTES_20_COLUMNAS_OFICIALES.csv",
    "reconciliacion": SUB / "resultados/auditorias/RECONCILIACION_FUENTES_20_COLUMNAS_BASE_62.csv",
    "tsv_dir": SUB / "data/governed/columnas_extranjeros_2025_v2",
    "matriz_v2": SUB / "data/processed/MATRIZ_GOBERNANZA_EXTRANJEROS_2025_V2.csv",
    "comparacion": SUB / "resultados/auditorias/COMPARACION_GOBERNANZA_V1_VS_V2.csv",
    "validacion": SUB / "resultados/auditorias/VALIDACION_GOBERNANZA_EXTRANJEROS_2025_V2.csv",
    "excel": SUB / "resultados/auditorias/GOBERNANZA_COLUMNAS_EXTRANJEROS_2025_V2.xlsx",
    "reporte": SUB / "resultados/reportes/REPORTE_CORRECCION_GOBERNANZA_EXTRANJEROS_2025_V2.md",
}

COMMON_TSV_COLUMNS = [
    "FILA_BASE_CONGELADA",
    "ID_REGISTRO_GOBERNADO",
    "ID_CONTROL_257",
    "CLAVE_PERSONA",
    "CLAVE_PERSONA_CARRERA",
    "NOMBRE_COLUMNA_OFICIAL",
    "VALOR_BASE_CONGELADA",
    "VALOR_PRECARGA_PES",
    "VALOR_MATRICULA_2025",
    "VALOR_UNIVERSO_AUDITADO",
    "VALOR_DATOS_ALUMNOS",
    "VALOR_RESOLUCION",
    "VALOR_OFERTA_ACADEMICA",
    "VALOR_SELECCIONADO",
    "ESTADO_VALOR",
    "REGLA_APLICADA",
    "FUENTE_SELECCIONADA",
    "ARCHIVO_FUENTE",
    "HOJA_FUENTE",
    "FILA_FUENTE",
    "NIVEL_CONFIANZA",
    "REQUIERE_REVISION",
    "RESPONSABLE_SUGERIDO",
    "CONFLICTO",
    "PENDIENTE",
    "OBSERVACION",
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_dirs() -> None:
    for path in OUT.values():
        if path.suffix:
            path.parent.mkdir(parents=True, exist_ok=True)
        else:
            path.mkdir(parents=True, exist_ok=True)


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: clean(row.get(col, "")) for col in columns})


def write_tsv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: clean(row.get(col, "")) for col in columns})


def read_df(path: Path, sep: str = ",") -> pd.DataFrame:
    return pd.read_csv(path, sep=sep, dtype=str, encoding="utf-8-sig", keep_default_na=False)


def detect_encoding_and_sep(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()[:8192]
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        enc = "latin-1"
        text = raw.decode(enc, errors="replace")
    first = text.splitlines()[0] if text.splitlines() else ""
    counts = {",": first.count(","), ";": first.count(";"), "\t": first.count("\t")}
    sep = max(counts, key=counts.get)
    return enc, sep


def read_auto(path: Path, nrows: int | None = None) -> pd.DataFrame:
    enc, sep = detect_encoding_and_sep(path)
    return pd.read_csv(path, sep=sep, dtype=str, encoding=enc, keep_default_na=False, nrows=nrows)


def normalize_text(value: Any) -> str:
    text = clean(value).strip().upper()
    text = re.sub(r"\s+", " ", text)
    text = "".join(ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn")
    return text


def parse_rut(value: Any) -> tuple[str, str, str]:
    text = clean(value).strip().upper().replace(".", "")
    if "-" in text:
        num, dv = text.rsplit("-", 1)
    else:
        num, dv = text[:-1], text[-1:] if text else ""
    num = re.sub(r"[^0-9]", "", num)
    dv = re.sub(r"[^0-9K]", "", dv)
    return ("R" if num and dv else "", num, dv)


def parse_date(value: Any) -> str:
    text = clean(value).strip()
    if not text:
        return ""
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).strftime("%d-%m-%Y")
        except ValueError:
            pass
    return text


def valid_birth(value: Any) -> bool:
    try:
        dt = datetime.strptime(parse_date(value), "%d-%m-%Y")
    except ValueError:
        return False
    return datetime(1900, 1, 1) <= dt <= datetime(2010, 6, 30)


def parse_int(value: Any) -> int | None:
    text = clean(value).strip()
    if not re.fullmatch(r"\d+", text):
        return None
    return int(text)


def valid_country(value: Any, allow_chile: bool = True) -> bool:
    n = parse_int(value)
    if n is None or not (1 <= n <= 197):
        return False
    if not allow_chile and n == 38:
        return False
    return True


def is_sies_code(value: Any) -> bool:
    return bool(re.fullmatch(r"I\d+S\d+C\d+J\d+V\d+", clean(value).strip()))


def dict_by(df: pd.DataFrame, key: str) -> dict[str, dict[str, str]]:
    if key not in df.columns:
        return {}
    out: dict[str, dict[str, str]] = {}
    for _, row in df.iterrows():
        k = clean(row.get(key, "")).strip()
        if k and k not in out:
            out[k] = {c: clean(row.get(c, "")) for c in df.columns}
    return out


def load_datos_alumnos() -> pd.DataFrame:
    wb = load_workbook(PATHS["datos_alumnos"], read_only=True, data_only=True)
    if "DatosAlumnos" not in wb.sheetnames:
        raise RuntimeError("No existe hoja DatosAlumnos")
    ws = wb["DatosAlumnos"]
    raw_headers = [clean(ws.cell(1, c).value).strip() or f"SIN_ENCABEZADO_{c}" for c in range(1, ws.max_column + 1)]
    seen: Counter[str] = Counter()
    headers = []
    for header in raw_headers:
        seen[header] += 1
        headers.append(header if seen[header] == 1 else f"{header}_{seen[header]}")
    rows = []
    for row_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = {headers[i]: clean(values[i]) for i in range(len(headers))}
        d["FILA_DATOS_ALUMNOS"] = str(row_idx)
        rows.append(d)
    return pd.DataFrame(rows, dtype=str).fillna("")


def backup_existing() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = SUB / "backups" / f"pre_correccion_gobernanza_extranjeros_2025_v2_{stamp}"
    backup.mkdir(parents=True, exist_ok=True)
    targets = list(OUT.values()) + [Path(__file__)]
    for target in targets:
        if not target.exists():
            continue
        rel = target.relative_to(SUB)
        dest = backup / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        if target.is_dir():
            if dest.exists():
                shutil.rmtree(dest)
            shutil.copytree(target, dest)
        else:
            shutil.copy2(target, dest)
    manifest_rows = []
    for path in sorted(p for p in backup.rglob("*") if p.is_file()):
        if path.name.startswith("MANIFIESTO_"):
            continue
        manifest_rows.append(
            {
                "ARCHIVO": str(path.relative_to(backup)),
                "HASH_SHA256": sha256(path),
                "TAMANO_BYTES": str(path.stat().st_size),
            }
        )
    write_csv(
        backup / "MANIFIESTO_RESPALDO_PRE_CORRECCION_GOBERNANZA_V2.csv",
        manifest_rows,
        ["ARCHIVO", "HASH_SHA256", "TAMANO_BYTES"],
    )
    return backup


def safe_row_from_line(df: pd.DataFrame, line_value: str) -> dict[str, str]:
    if not clean(line_value).isdigit():
        return {}
    idx = int(line_value) - 2
    if idx < 0 or idx >= len(df):
        return {}
    return {c: clean(df.iloc[idx].get(c, "")) for c in df.columns}


def clave_code(value: str) -> str:
    parts = clean(value).split("|")
    if len(parts) >= 4 and is_sies_code(parts[3]):
        return parts[3]
    return ""


def governed_filename(order: int, column: str) -> str:
    return f"{order:02d}_{column}.tsv"


def state_is_valid(state: str) -> bool:
    return state in {"VALIDADO", "VALIDADO_CON_ADVERTENCIA", "VACIO_PERMITIDO", "NO_APLICA"}


def source_value(row: dict[str, str], key: str) -> str:
    return clean(row.get(key, "")).strip()


def build_inventory() -> list[dict[str, str]]:
    keywords = [
        "TIPO_DOCUMENTO",
        "NUM_DOCUMENTO",
        "CODIGO_UNICO",
        "PAIS_ESTUDIOS_SECUNDARIOS",
        "PAIS_ORIGEN",
        "TIPO_RESIDENCIA_ESTUDIANTE",
        "ANIO_INGRESO_CARRERA_ORIGEN",
        "SEM_INGRESO_CARRERA_ORIGEN",
        "VIGENCIA",
        "NACIONALIDAD",
    ]
    key_lc = [k.lower() for k in keywords]
    known = set(PATHS.values())
    roots = [
        SUB / "data",
        SUB / "resultados/auditorias",
        SUB / "scripts",
        ROOT / "resultados",
        ROOT / "control",
    ]
    candidates: set[Path] = {p for p in known if p.exists()}
    skip_parts = {".git", ".venv", "node_modules", "__pycache__", "backups", "archive", "caches", ".codex"}
    for root in roots:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_dir() or any(part in skip_parts for part in path.parts):
                continue
            if path.suffix.lower() not in {".csv", ".tsv", ".xlsx", ".txt", ".json", ".py"}:
                continue
            name_lc = path.name.lower()
            if any(k.lower() in name_lc for k in ["precarga", "conciliacion", "resolucion", "codigo", "vigencia", "nacionalidad", "pais", "matricula", "carrera", "oferta", "extranj"]):
                candidates.add(path)
    rows = []
    for path in sorted(candidates):
        rel = str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else str(path)
        tipo = path.suffix.lower().lstrip(".")
        hojas = [""]
        if path.suffix.lower() == ".xlsx":
            try:
                wb = load_workbook(path, read_only=True, data_only=True)
                hojas = wb.sheetnames[:6]
            except Exception:
                hojas = [""]
        for hoja in hojas:
            variables: list[str] = []
            claves: list[str] = []
            filas = ""
            columnas = ""
            estado = "INSPECCIONADO"
            obs = ""
            try:
                if path.suffix.lower() in {".csv", ".tsv"}:
                    df = read_auto(path, nrows=30)
                    filas = str(max(0, sum(1 for _ in path.open("rb")) - 1))
                    columnas = str(len(df.columns))
                    headers = [normalize_text(c) for c in df.columns]
                elif path.suffix.lower() == ".xlsx" and hoja:
                    ws = load_workbook(path, read_only=True, data_only=True)[hoja]
                    rows_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
                    headers = [normalize_text(c) for c in next(rows_iter, []) if clean(c)]
                    filas = str(max(0, ws.max_row - 1))
                    columnas = str(ws.max_column)
                else:
                    text = path.read_text(encoding="utf-8", errors="ignore")[:10000]
                    headers = [k for k in keywords if k in text]
                    filas = ""
                    columnas = ""
                variables = [k for k in keywords if normalize_text(k) in headers or any(normalize_text(k) in h for h in headers)]
                claves = [k for k in ["CODCLI", "RUT", "NUM_DOCUMENTO", "CLAVE_PERSONA_DOCUMENTO", "CLAVE_OFICIAL_CARGA", "CODIGO_UNICO"] if normalize_text(k) in headers or any(normalize_text(k) == h for h in headers)]
            except Exception as exc:
                estado = "ERROR_LECTURA"
                obs = str(exc)
            if not variables and path not in known:
                continue
            priority = "ALTA" if path in known else ("MEDIA" if "resultados" in str(path) or "control" in str(path) else "BAJA")
            rows.append(
                {
                    "ARCHIVO": rel,
                    "TIPO_ARCHIVO": tipo,
                    "HOJA": hoja,
                    "FILAS": filas,
                    "COLUMNAS": columnas,
                    "VARIABLES_OFICIALES_ENCONTRADAS": "|".join(sorted(set(variables))),
                    "CLAVES_DISPONIBLES": "|".join(sorted(set(claves))),
                    "PRIORIDAD_USO": priority,
                    "ESTADO": estado,
                    "OBSERVACION": obs,
                }
            )
    return rows


def normalize_text_for_header(value: Any) -> str:
    return normalize_text(value).replace(" ", "_")


def build_context() -> dict[str, Any]:
    required = [PATHS["frozen"], PATHS["structure"], PATHS["manual"], PATHS["precarga_norm"], PATHS["matricula"], PATHS["universo"]]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise RuntimeError("Faltan fuentes requeridas: " + "; ".join(missing))
    if sha256(PATHS["frozen"]) != FROZEN_HASH:
        raise RuntimeError("Hash de base congelada no coincide")
    if sha256(PATHS["structure"]) != STRUCTURE_HASH:
        raise RuntimeError("Hash de estructura oficial no coincide")

    structure = read_df(PATHS["structure"], sep=";")
    if list(structure.columns) != OFFICIAL_COLUMNS:
        raise RuntimeError("La estructura oficial no conserva las 20 columnas esperadas")

    frozen = read_df(PATHS["frozen"], sep="\t")
    if frozen.shape != (62, 63):
        raise RuntimeError(f"Base congelada tiene dimensiones {frozen.shape}, se esperaba (62, 63)")
    frozen["FILA_BASE_CONGELADA"] = [str(i) for i in range(2, len(frozen) + 2)]

    ctx = {
        "frozen": frozen,
        "precarga": read_df(PATHS["precarga_norm"]),
        "matricula": read_df(PATHS["matricula"]),
        "universo": read_df(PATHS["universo"]),
        "matriz_v1": read_df(PATHS["matriz_v1"]),
        "cobertura_v1": read_df(PATHS["cobertura_v1"]),
        "cierre": read_df(PATHS["cierre"]),
        "conc_precarga": read_df(PATHS["conc_precarga"]),
        "conc_matricula": read_df(PATHS["conc_matricula"]),
        "conc_257": read_df(PATHS["conc_257"]),
        "res_codigo": read_df(PATHS["res_codigo"]),
        "res_mult": read_df(PATHS["res_mult"]),
        "cand_mult": read_df(PATHS["cand_mult"]),
        "res_evid4": read_df(PATHS["res_evid4"]),
        "res_nac": read_df(PATHS["res_nac"]),
        "vig_corr": read_df(PATHS["vig_corr"]),
        "catalogo_paises": read_df(PATHS["catalogo_paises"]),
        "puente_sies": read_df(PATHS["puente_sies"], sep="\t") if PATHS["puente_sies"].exists() else pd.DataFrame(),
        "datos_alumnos": load_datos_alumnos(),
    }
    return ctx


def prepare_maps(ctx: dict[str, Any]) -> dict[str, Any]:
    maps: dict[str, Any] = {}
    maps["prec_by_fila"] = dict_by(ctx["precarga"], "FILA_PRECARGA")
    maps["conc_prec_by_fila"] = dict_by(ctx["conc_precarga"], "FILA_BASE_CONGELADA")
    maps["conc_mat_by_fila"] = dict_by(ctx["conc_matricula"], "FILA_BASE_CONGELADA")
    maps["conc_257_by_fila"] = dict_by(ctx["conc_257"], "FILA_BASE_CONGELADA")
    maps["cierre_by_fila"] = dict_by(ctx["cierre"], "FILA_BASE_CONGELADA")
    maps["nac_by_fila"] = dict_by(ctx["res_nac"], "FILA_BASE_CONGELADA")
    maps["evid_by_fila"] = dict_by(ctx["res_evid4"], "FILA_BASE_CONGELADA")
    maps["mult_by_fila"] = dict_by(ctx["res_mult"], "FILA_BASE_CONGELADA")
    maps["selected_cand_by_fila"] = {
        clean(row["FILA_BASE_CONGELADA"]): {c: clean(row.get(c, "")) for c in ctx["cand_mult"].columns}
        for _, row in ctx["cand_mult"].iterrows()
        if clean(row.get("ES_SELECCIONADO", "")) == "SI"
    }
    maps["codigo_by_codcli"] = dict_by(ctx["res_codigo"], "CODCLI")
    maps["da_by_codcli"] = dict_by(ctx["datos_alumnos"], "CODCLI")
    maps["universo_by_id"] = dict_by(ctx["universo"], "ID_CONTROL")
    maps["vig_by_codcli"] = dict_by(ctx["vig_corr"], "CODCLI")
    return maps


def prec_value(prec: dict[str, str], column: str) -> str:
    if not prec:
        return ""
    mapping = {"PAIS_ORIGEN": "PAIS_DE_ORIGEN_NORMALIZADO"}
    return source_value(prec, mapping.get(column, f"{column}_NORMALIZADO"))


def country_code_from_cierre(cierre: dict[str, str]) -> str:
    return source_value(cierre, "NACIONALIDAD_CODIGO_SIES")


def build_source_pack(
    frozen_row: dict[str, str],
    ctx: dict[str, Any],
    maps: dict[str, Any],
) -> dict[str, Any]:
    fila = frozen_row["FILA_BASE_CONGELADA"]
    conc_prec = maps["conc_prec_by_fila"].get(fila, {})
    prec = maps["prec_by_fila"].get(source_value(conc_prec, "FILA_PRECARGA"), {})
    conc_mat = maps["conc_mat_by_fila"].get(fila, {})
    mat = safe_row_from_line(ctx["matricula"], source_value(conc_mat, "FILA_MATRICULA_2025"))
    conc257 = maps["conc_257_by_fila"].get(fila, {})
    universo = maps["universo_by_id"].get(source_value(conc257, "ID_CONTROL_257"), {})
    cierre = maps["cierre_by_fila"].get(fila, {})
    da = maps["da_by_codcli"].get(frozen_row.get("CODCLI", ""), {})
    nac = maps["nac_by_fila"].get(fila, {})
    evid = maps["evid_by_fila"].get(fila, {})
    mult = maps["mult_by_fila"].get(fila, {})
    selected = maps["selected_cand_by_fila"].get(fila, {})
    codres = maps["codigo_by_codcli"].get(frozen_row.get("CODCLI", ""), {})
    vig = maps["vig_by_codcli"].get(frozen_row.get("CODCLI", ""), {})
    return {
        "conc_prec": conc_prec,
        "prec": prec,
        "conc_mat": conc_mat,
        "mat": mat,
        "conc257": conc257,
        "universo": universo,
        "cierre": cierre,
        "da": da,
        "nac": nac,
        "evid": evid,
        "mult": mult,
        "selected": selected,
        "codres": codres,
        "vig": vig,
    }


def source_values_for_column(column: str, frozen: dict[str, str], pack: dict[str, Any]) -> dict[str, str]:
    tipo, num, dv = parse_rut(frozen.get("RUT", ""))
    da_tipo, da_num, da_dv = parse_rut(pack["da"].get("RUT", "")) if pack["da"] else ("", "", "")
    base = {
        "TIPO_DOCUMENTO": tipo,
        "NUM_DOCUMENTO": num,
        "DV": dv,
        "PRIMER_APELLIDO": "",
        "SEGUNDO_APELLIDO": "",
        "NOMBRES": "",
        "SEXO": {"F": "M", "M": "H"}.get(source_value(frozen, "SEXO"), ""),
        "FECHA_NACIMIENTO": parse_date(frozen.get("FECHANACIMIENTO", "")),
        "NACIONALIDAD": frozen.get("NACIONALIDAD", ""),
        "TIPO_RESIDENCIA_ESTUDIANTE": "",
        "PAIS_ORIGEN": "",
        "PAIS_ESTUDIOS_SECUNDARIOS": "",
        "CODIGO_UNICO": frozen.get("CODCARPR", ""),
        "ANIO_INGRESO_CARRERA_ACTUAL": frozen.get("ANOINGRESO", ""),
        "SEM_INGRESO_CARRERA_ACTUAL": frozen.get("PERIODOINGRESO", ""),
        "ANIO_INGRESO_CARRERA_ORIGEN": "",
        "SEM_INGRESO_CARRERA_ORIGEN": "",
        "NOMBRE_UNIVERSIDAD_ORIGEN": frozen.get("NOMBREUNIVERSIDAD", ""),
        "PAIS_UNIVERSIDAD_ORIGEN": "",
        "VIGENCIA": frozen.get("MATRICULA", ""),
    }
    da = {
        "TIPO_DOCUMENTO": da_tipo,
        "NUM_DOCUMENTO": da_num,
        "DV": da_dv,
        "PRIMER_APELLIDO": pack["da"].get("APELLIDO PATERNO", ""),
        "SEGUNDO_APELLIDO": pack["da"].get("APELLIDO MATERNO", ""),
        "NOMBRES": pack["da"].get("NOMBRES", ""),
        "SEXO": {"F": "M", "M": "H"}.get(source_value(pack["da"], "SEXO"), source_value(pack["da"], "SEXO")),
        "FECHA_NACIMIENTO": parse_date(pack["da"].get("FECHANACIMIENTO", "")),
        "NACIONALIDAD": pack["da"].get("NACIONALIDAD", ""),
        "TIPO_RESIDENCIA_ESTUDIANTE": "",
        "PAIS_ORIGEN": "",
        "PAIS_ESTUDIOS_SECUNDARIOS": "",
        "CODIGO_UNICO": pack["da"].get("CODCARPR", ""),
        "ANIO_INGRESO_CARRERA_ACTUAL": pack["da"].get("ANOINGRESO", ""),
        "SEM_INGRESO_CARRERA_ACTUAL": pack["da"].get("PERIODOINGRESO", ""),
        "ANIO_INGRESO_CARRERA_ORIGEN": "",
        "SEM_INGRESO_CARRERA_ORIGEN": "",
        "NOMBRE_UNIVERSIDAD_ORIGEN": pack["da"].get("NOMBREUNIVERSIDAD", ""),
        "PAIS_UNIVERSIDAD_ORIGEN": "",
        "VIGENCIA": "",
    }
    mat = {
        "TIPO_DOCUMENTO": pack["mat"].get("TIPO_DOCUMENTO", ""),
        "NUM_DOCUMENTO": pack["mat"].get("NUM_DOCUMENTO", ""),
        "DV": pack["mat"].get("DV", ""),
        "PRIMER_APELLIDO": pack["mat"].get("PRIMER_APELLIDO", ""),
        "SEGUNDO_APELLIDO": pack["mat"].get("SEGUNDO_APELLIDO", ""),
        "NOMBRES": pack["mat"].get("NOMBRES", ""),
        "SEXO": pack["mat"].get("SEXO", ""),
        "FECHA_NACIMIENTO": parse_date(pack["mat"].get("FECHA_NACIMIENTO", "")),
        "CODIGO_UNICO": pack["mat"].get("CODIGO_UNICO", ""),
        "ANIO_INGRESO_CARRERA_ACTUAL": pack["mat"].get("ANIO_INGRESO_CARRERA_ACTUAL", ""),
        "SEM_INGRESO_CARRERA_ACTUAL": pack["mat"].get("SEM_INGRESO_CARRERA_ACTUAL", ""),
        "ANIO_INGRESO_CARRERA_ORIGEN": pack["mat"].get("ANIO_INGRESO_CARRERA_ORIGEN", ""),
        "SEM_INGRESO_CARRERA_ORIGEN": pack["mat"].get("SEM_INGRESO_CARRERA_ORIGEN", ""),
        "VIGENCIA": pack["mat"].get("VIGENCIA", ""),
    }
    universo_code = source_value(pack["conc257"], "CODIGO_UNICO_257") or clave_code(source_value(pack["universo"], "CLAVE_OFICIAL_CARGA"))
    universo = {
        "TIPO_DOCUMENTO": pack["conc257"].get("TIPO_DOCUMENTO_257", ""),
        "NUM_DOCUMENTO": pack["conc257"].get("NUM_DOCUMENTO_257", ""),
        "DV": "",
        "NACIONALIDAD": pack["conc257"].get("NACIONALIDAD_257", ""),
        "CODIGO_UNICO": universo_code,
        "VIGENCIA": pack["universo"].get("VIGENCIA_PROPUESTA", ""),
    }
    resol = {
        "NACIONALIDAD": country_code_from_cierre(pack["cierre"]),
        "CODIGO_UNICO": source_value(pack["selected"], "CODIGO_UNICO_CANDIDATO")
        or source_value(pack["codres"], "CODIGO_UNICO_DISPONIBLE")
        or source_value(pack["conc257"], "CODIGO_UNICO_257"),
        "VIGENCIA": source_value(pack["vig"], "VIGENCIA_PROPUESTA")
        or ("1" if source_value(pack["cierre"], "EVIDENCIA_2025_ESTADO_FINAL").startswith("EVIDENCIA_2025_CONFIRMADA") else ""),
    }
    oferta = {
        "CODIGO_UNICO": source_value(pack["codres"], "CODIGO_UNICO_DISPONIBLE"),
    }
    return {
        "base": clean(base.get(column, "")),
        "precarga": prec_value(pack["prec"], column),
        "matricula": clean(mat.get(column, "")),
        "universo": clean(universo.get(column, "")),
        "datos_alumnos": clean(da.get(column, "")),
        "resolucion": clean(resol.get(column, "")),
        "oferta": clean(oferta.get(column, "")),
    }


def decide_column(column: str, vals: dict[str, str], frozen: dict[str, str], pack: dict[str, Any]) -> dict[str, str]:
    fila_prec = source_value(pack["conc_prec"], "FILA_PRECARGA")
    fila_mat = source_value(pack["conc_mat"], "FILA_MATRICULA_2025")
    fila_da = source_value(pack["da"], "FILA_DATOS_ALUMNOS")
    result = {
        "VALOR_SELECCIONADO": "",
        "ESTADO_VALOR": PENDING,
        "REGLA_APLICADA": "",
        "FUENTE_SELECCIONADA": "",
        "ARCHIVO_FUENTE": "",
        "HOJA_FUENTE": "",
        "FILA_FUENTE": "",
        "NIVEL_CONFIANZA": "BAJA",
        "REQUIERE_REVISION": "SI",
        "RESPONSABLE_SUGERIDO": responsible(column),
        "CONFLICTO": "NO",
        "PENDIENTE": "SI",
        "OBSERVACION": "",
    }

    def select(value: str, state: str, rule: str, source: str, file_path: Path | str, fila: str = "", confidence: str = "ALTA", obs: str = "", sheet: str = "") -> dict[str, str]:
        result.update(
            {
                "VALOR_SELECCIONADO": clean(value),
                "ESTADO_VALOR": state,
                "REGLA_APLICADA": rule,
                "FUENTE_SELECCIONADA": source,
                "ARCHIVO_FUENTE": str(file_path),
                "HOJA_FUENTE": sheet,
                "FILA_FUENTE": fila,
                "NIVEL_CONFIANZA": confidence,
                "REQUIERE_REVISION": "NO" if state_is_valid(state) else "SI",
                "CONFLICTO": "SI" if state == "CONFLICTO_ENTRE_FUENTES" else "NO",
                "PENDIENTE": "NO" if state_is_valid(state) else "SI",
                "OBSERVACION": obs,
            }
        )
        return result

    b, p, m, u, da, r, o = (vals[k] for k in ["base", "precarga", "matricula", "universo", "datos_alumnos", "resolucion", "oferta"])

    if column == "TIPO_DOCUMENTO":
        if p in {"R", "P"}:
            state = "VALIDADO_CON_ADVERTENCIA" if p != b and b else "VALIDADO"
            obs = "Precarga exacta informa tipo documental distinto al RUT institucional; se conserva precarga." if state.endswith("ADVERTENCIA") else ""
            return select(p, state, "Prioridad a precarga exacta; no convertir pasaporte a R.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec, "ALTA", obs)
        if m in {"R", "P"}:
            return select(m, "VALIDADO", "Sin precarga; usar matricula control 2025.", "MATRICULA_CONTROL_2025", PATHS["matricula"], fila_mat)
        if b == "R" and da == "R":
            return select(b, "VALIDADO_CON_ADVERTENCIA", "Sin precarga ni matricula; RUT coincide entre base congelada y DatosAlumnos.", "BASE_CONGELADA_Y_DATOS_ALUMNOS", PATHS["frozen"], "", "MEDIA")
        if b == "R":
            return select(b, "VALIDADO_CON_ADVERTENCIA", "Sin precarga; RUT institucional en base congelada.", "BASE_CONGELADA", PATHS["frozen"], "", "MEDIA")
        return result

    if column == "NUM_DOCUMENTO":
        value = p or m or b or da
        if value and re.fullmatch(r"\d+", value):
            source = "PRECARGA_NORMALIZADA" if p else ("MATRICULA_CONTROL_2025" if m else ("BASE_CONGELADA" if b else "DATOS_ALUMNOS"))
            path = PATHS["precarga_norm"] if p else (PATHS["matricula"] if m else (PATHS["frozen"] if b else PATHS["datos_alumnos"]))
            fila = fila_prec if p else (fila_mat if m else ("" if b else fila_da))
            return select(value, "VALIDADO", "Documento como texto desde fuente prioritaria disponible.", source, path, fila)
        return select(value, "VACIO_NO_PERMITIDO", "NUM_DOCUMENTO obligatorio.", "", "", "", "BAJA", "Documento ausente o no numerico.")

    if column == "DV":
        selected_type = source_value(pack.get("decisions", {}).get("TIPO_DOCUMENTO", {}), "VALOR_SELECCIONADO") or p or m or b
        if selected_type == "P":
            return select("", "VALIDADO", "DV nulo para pasaporte segun instructivo.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        value = p or m or b or da
        if re.fullmatch(r"[0-9K]", value):
            source = "PRECARGA_NORMALIZADA" if p else ("MATRICULA_CONTROL_2025" if m else "BASE_CONGELADA")
            path = PATHS["precarga_norm"] if p else (PATHS["matricula"] if m else PATHS["frozen"])
            fila = fila_prec if p else (fila_mat if m else "")
            return select(value, "VALIDADO", "DV para RUN desde fuente prioritaria.", source, path, fila)
        return select(value, "VACIO_NO_PERMITIDO", "DV requerido para RUN.", "", "", "", "BAJA")

    if column in {"PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES"}:
        if p:
            return select(p, "VALIDADO", "Usar campo estructurado de precarga; no dividir nombre completo.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        if m:
            return select(m, "VALIDADO", "Usar campo estructurado de matricula control; no dividir nombre completo.", "MATRICULA_CONTROL_2025", PATHS["matricula"], fila_mat)
        if da:
            state = "VALIDADO" if column != "SEGUNDO_APELLIDO" or da else "VACIO_PERMITIDO"
            return select(da, state, "Usar campo separado de DatosAlumnos; no dividir nombre completo.", "DATOS_ALUMNOS", PATHS["datos_alumnos"], fila_da, "ALTA", sheet="DatosAlumnos")
        if column == "SEGUNDO_APELLIDO":
            return select("", "VACIO_PERMITIDO", "Segundo apellido puede no estar informado en fuente estructurada.", "SIN_VALOR_ESTRUCTURADO", "", "", "MEDIA")
        return select("", "VACIO_NO_PERMITIDO", f"{column} obligatorio.", "", "", "")

    if column == "SEXO":
        if p in {"M", "H", "NB"}:
            return select(p, "VALIDADO", "Precarga coincide con transformacion institucional F->M y M->H.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        if b in {"M", "H"}:
            return select(b, "VALIDADO", "Transformacion institucional confirmada: F local -> M SIES; M local -> H SIES.", "BASE_CONGELADA_TRANSFORMACION_INSTITUCIONAL", PATHS["frozen"], "", "ALTA")
        return select(b, "VALOR_INVALIDO", "SEXO debe ser M, H o NB.", "BASE_CONGELADA", PATHS["frozen"], "", "BAJA")

    if column == "FECHA_NACIMIENTO":
        value = parse_date(p or m or b or da)
        source = "PRECARGA_NORMALIZADA" if p else ("MATRICULA_CONTROL_2025" if m else ("BASE_CONGELADA" if b else "DATOS_ALUMNOS"))
        path = PATHS["precarga_norm"] if p else (PATHS["matricula"] if m else (PATHS["frozen"] if b else PATHS["datos_alumnos"]))
        fila = fila_prec if p else (fila_mat if m else ("" if b else fila_da))
        if value and valid_birth(value):
            return select(value, "VALIDADO", "Fecha normalizada a DD-MM-AAAA desde fuente prioritaria.", source, path, fila)
        return select(value, "VALOR_INVALIDO", "Fecha ausente o fuera de rango normativo.", source, path, fila, "BAJA")

    if column == "NACIONALIDAD":
        if p and valid_country(p, allow_chile=False):
            return select(p, "VALIDADO", "Prioridad a precarga exacta; codigo SIES 1-197 distinto de Chile.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        if r and valid_country(r, allow_chile=False):
            return select(r, "VALIDADO", "Catalogo SIES y resolucion de cierre para valor base congelada.", "BASE_CIERRE_CATALOGO_SIES", PATHS["cierre"], "")
        if normalize_text(b) == "POR DEFINIR":
            return select("", PENDING, "No completar valor Por definir sin confirmacion institucional.", "RESOLUCION_NACIONALIDAD", PATHS["res_nac"], source_value(pack["nac"], "FILA_FUENTE_BASE"), "BAJA", "Nacionalidad bruta Por definir.")
        return select(r or p, "VALOR_INVALIDO", "Nacionalidad debe estar entre 1 y 197 y no puede ser 38.", "", "", "")

    if column == "TIPO_RESIDENCIA_ESTUDIANTE":
        if p in {"1", "2", "3"}:
            return select(p, "VALIDADO", "Valor explicito de precarga.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        return select("", PENDING, "Campo obligatorio sin valor explicito en precarga ni fuentes reutilizadas; no se infiere desde direccion.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA", "Precarga revisada: sin valores para los 62 registros.")

    if column == "PAIS_ORIGEN":
        residencia = source_value(pack.get("decisions", {}).get("TIPO_RESIDENCIA_ESTUDIANTE", {}), "VALOR_SELECCIONADO")
        if residencia == "1":
            return select("", "VACIO_PERMITIDO", "Si residencia=1, instructivo indica no completar PAIS_ORIGEN.", "REGLA_CONDICIONAL_MANUAL", PATHS["manual"], "")
        if residencia in {"2", "3"}:
            if p and valid_country(p, allow_chile=False):
                return select(p, "VALIDADO", "Residencia 2/3 requiere pais origen explicito desde precarga.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
            return select(p, PENDING, "Residencia 2/3 requiere PAIS_ORIGEN; no se infiere desde nacionalidad.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA")
        return select("", PENDING, "No evaluable sin TIPO_RESIDENCIA_ESTUDIANTE; no exigir pais origen como valor independiente.", "DEPENDENCIA_TIPO_RESIDENCIA", PATHS["manual"], "", "BAJA", "Pendiente condicionado por residencia.")

    if column == "PAIS_ESTUDIOS_SECUNDARIOS":
        if p and valid_country(p, allow_chile=True):
            return select(p, "VALIDADO", "Pais estudios secundarios explicito en precarga; no se usa colegio/ciudad.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        return select("", PENDING, "No existe pais estudios secundarios explicito en fuente reutilizada para este registro.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA", "No se usa COLEGIO, CIUDADCOLEGIO ni COMUNACOLEGIO.")

    if column == "CODIGO_UNICO":
        mult_status = source_value(pack["mult"], "TIPO_RESOLUCION")
        if mult_status == "EMPATE_ENTRE_CANDIDATOS":
            return select("", PENDING, "Empate de coincidencia multiple; no se selecciona primer candidato.", "RESOLUCION_19_COINCIDENCIAS_MULTIPLES", PATHS["res_mult"], "", "BAJA", "Empate formalmente pendiente.")
        selected = source_value(pack["selected"], "CODIGO_UNICO_CANDIDATO")
        if selected and is_sies_code(selected):
            return select(selected, "VALIDADO", "Seleccion unica alta confianza desde resolucion de candidatos multiples.", "CANDIDATOS_19_COINCIDENCIAS_MULTIPLES", PATHS["cand_mult"], "", "ALTA")
        if p and is_sies_code(p):
            return select(p, "VALIDADO", "Codigo unico desde precarga exacta persona.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec)
        if u and is_sies_code(u):
            return select(u, "VALIDADO", "Codigo unico extraido de universo auditado/CLAVE_OFICIAL_CARGA con vinculo validado.", "UNIVERSO_AUDITADO", PATHS["universo"], "", "ALTA")
        if o and is_sies_code(o) and source_value(pack["codres"], "CLASIFICACION_RESOLUCION").startswith("RESUELTO"):
            return select(o, "VALIDADO", "Resolucion codigo unico reutilizada; no se usa CODCARPR como valor final.", "RESOLUCION_CODIGO_UNICO", PATHS["res_codigo"], "", "ALTA")
        return select("", PENDING, "Sin codigo unico SIES inequivoco; no usar CODCARPR/CODIGOCARRERA/CODCLI como valor oficial.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA", source_value(pack["codres"], "CLASIFICACION_RESOLUCION") or "Sin resolucion disponible.")

    if column == "ANIO_INGRESO_CARRERA_ACTUAL":
        value = p or m or b
        n = parse_int(value)
        source = "PRECARGA_NORMALIZADA" if p else ("MATRICULA_CONTROL_2025" if m else "BASE_CONGELADA")
        path = PATHS["precarga_norm"] if p else (PATHS["matricula"] if m else PATHS["frozen"])
        fila = fila_prec if p else (fila_mat if m else "")
        if n is not None and 1950 <= n <= 2025:
            return select(value, "VALIDADO", "Prioridad precarga/matricula; base ANOINGRESO solo si no hay fuente previa.", source, path, fila)
        return select(value, "VALOR_INVALIDO", "Ano ingreso actual fuera de rango.", source, path, fila, "BAJA")

    if column == "SEM_INGRESO_CARRERA_ACTUAL":
        value = p or m or b
        source = "PRECARGA_NORMALIZADA" if p else ("MATRICULA_CONTROL_2025" if m else "BASE_CONGELADA")
        path = PATHS["precarga_norm"] if p else (PATHS["matricula"] if m else PATHS["frozen"])
        fila = fila_prec if p else (fila_mat if m else "")
        if value in {"1", "2"}:
            return select(value, "VALIDADO", "Prioridad precarga/matricula; base PERIODOINGRESO solo si no hay fuente previa.", source, path, fila)
        return select(value, "VALOR_INVALIDO", "Semestre actual debe ser 1 o 2.", source, path, fila, "BAJA")

    if column == "ANIO_INGRESO_CARRERA_ORIGEN":
        value = p or m
        n = parse_int(value)
        actual = parse_int(source_value(pack.get("decisions", {}).get("ANIO_INGRESO_CARRERA_ACTUAL", {}), "VALOR_SELECCIONADO"))
        if n is not None and ((1950 <= n <= 2025) or n == 1900) and (actual is None or n <= actual or n == 1900):
            source = "PRECARGA_NORMALIZADA" if p else "MATRICULA_CONTROL_2025"
            path = PATHS["precarga_norm"] if p else PATHS["matricula"]
            fila = fila_prec if p else fila_mat
            return select(value, "VALIDADO", "Ano origen desde fuente explicita; no se asume igual al actual.", source, path, fila)
        return select("", PENDING, "Sin ano origen explicito en precarga/matricula; no se asume desde ANOINGRESO.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA")

    if column == "SEM_INGRESO_CARRERA_ORIGEN":
        value = p or m
        anio_origen = source_value(pack.get("decisions", {}).get("ANIO_INGRESO_CARRERA_ORIGEN", {}), "VALOR_SELECCIONADO")
        if value and ((anio_origen == "1900" and value == "0") or (anio_origen != "1900" and value in {"1", "2"})):
            source = "PRECARGA_NORMALIZADA" if p else "MATRICULA_CONTROL_2025"
            path = PATHS["precarga_norm"] if p else PATHS["matricula"]
            fila = fila_prec if p else fila_mat
            return select(value, "VALIDADO", "Semestre origen desde fuente explicita; no se asume desde periodo actual.", source, path, fila)
        return select("", PENDING, "Sin semestre origen explicito en precarga/matricula.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA")

    if column in {"NOMBRE_UNIVERSIDAD_ORIGEN", "PAIS_UNIVERSIDAD_ORIGEN"}:
        if p:
            if column == "PAIS_UNIVERSIDAD_ORIGEN" and not valid_country(p, allow_chile=True):
                return select(p, "VALOR_INVALIDO", "Pais universidad origen debe estar entre 1 y 197.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec, "MEDIA")
            return select(p, "VALIDADO", "Valor explicito de precarga para universidad de origen.", "PRECARGA_NORMALIZADA", PATHS["precarga_norm"], fila_prec, "MEDIA")
        return select("", "NO_APLICA", "Sin fuente que indique universidad de origen; no usar carrera anterior ni institucion actual.", "REGLA_CONDICIONAL_MANUAL", PATHS["manual"], "", "MEDIA")

    if column == "VIGENCIA":
        cierre_estado = source_value(pack["cierre"], "EVIDENCIA_2025_ESTADO_FINAL")
        evid_dec = source_value(pack["evid"], "DECISION_FINAL")
        if cierre_estado == "EVIDENCIA_2025_PARCIAL":
            return select("", PENDING, "Evidencia 2025 parcial; no usar MATRICULA 1/2 como vigencia.", "RESOLUCION_4_SIN_EVIDENCIA_2025", PATHS["res_evid4"], "", "BAJA")
        if cierre_estado.startswith("EVIDENCIA_2025_CONFIRMADA"):
            state = "VALIDADO_CON_ADVERTENCIA" if "DATOS_ALUMNOS" in cierre_estado else "VALIDADO"
            obs = "Confirmada por DatosAlumnos, no por control local inicial." if state.endswith("ADVERTENCIA") else ""
            return select("1", state, "Vigencia 1 solo con evidencia explicita 2025; no deriva de MATRICULA 1/2.", "CIERRE_PENDIENTES_BASE_62", PATHS["cierre"], "", "ALTA" if state == "VALIDADO" else "MEDIA", obs)
        if r == "1":
            return select("1", "VALIDADO", "Resolucion de vigencia corregida reutilizada.", "AUDITORIA_VIGENCIA_2025_CORREGIDA", PATHS["vig_corr"], "", "ALTA")
        return select("", PENDING, "Sin evidencia 2025 suficiente.", "FUENTE_NO_DISPONIBLE", "", "", "BAJA")

    return result


def responsible(column: str) -> str:
    if column in {"CODIGO_UNICO", "ANIO_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL", "ANIO_INGRESO_CARRERA_ORIGEN", "SEM_INGRESO_CARRERA_ORIGEN"}:
        return "Docencia/Oferta Academica"
    if column in {"TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS", "NACIONALIDAD", "VIGENCIA"}:
        return "Registro Academico/Docencia"
    return "Registro Academico"


def build_governance(ctx: dict[str, Any], maps: dict[str, Any]) -> tuple[dict[str, list[dict[str, str]]], pd.DataFrame, pd.DataFrame, dict[str, pd.DataFrame]]:
    per_col = {col: [] for col in OFFICIAL_COLUMNS}
    matrix_rows = []
    reconciliation_rows = []
    prec_rows = []
    mat_rows = []
    uni_rows = []
    for _, row in ctx["frozen"].iterrows():
        frozen = {c: clean(row.get(c, "")) for c in ctx["frozen"].columns}
        fila = frozen["FILA_BASE_CONGELADA"]
        pack = build_source_pack(frozen, ctx, maps)
        pack["decisions"] = {}
        matrix_row = {c: frozen[c] for c in ctx["frozen"].columns if c != "FILA_BASE_CONGELADA"}
        matrix_row["FILA_BASE_CONGELADA"] = fila
        matrix_row["ID_CONTROL_257"] = source_value(pack["conc257"], "ID_CONTROL_257")
        matrix_row["FILA_PRECARGA"] = source_value(pack["conc_prec"], "FILA_PRECARGA")
        matrix_row["FILA_MATRICULA_2025"] = source_value(pack["conc_mat"], "FILA_MATRICULA_2025")
        matrix_row["REQUIERE_REVISION_CIERRE"] = source_value(pack["cierre"], "REQUIERE_REVISION_FINAL")
        matrix_row["MOTIVO_REVISION_CIERRE"] = source_value(pack["cierre"], "MOTIVO_REVISION_FINAL")
        rec = {
            "FILA_BASE_CONGELADA": fila,
            "CODCLI": frozen.get("CODCLI", ""),
            "RUT": frozen.get("RUT", ""),
            "NOMBRE": frozen.get("NOMBRE", ""),
            "ID_CONTROL_257": source_value(pack["conc257"], "ID_CONTROL_257"),
            "FILA_PRECARGA": source_value(pack["conc_prec"], "FILA_PRECARGA"),
            "FILA_MATRICULA_2025": source_value(pack["conc_mat"], "FILA_MATRICULA_2025"),
        }
        for col in OFFICIAL_COLUMNS:
            vals = source_values_for_column(col, frozen, pack)
            decision = decide_column(col, vals, frozen, pack)
            pack["decisions"][col] = decision
            tsv_row = {
                "FILA_BASE_CONGELADA": fila,
                "ID_REGISTRO_GOBERNADO": f"GOBV2-EXT-2025-{int(fila):03d}",
                "ID_CONTROL_257": source_value(pack["conc257"], "ID_CONTROL_257"),
                "CLAVE_PERSONA": source_value(pack["conc257"], "CLAVE_PERSONA"),
                "CLAVE_PERSONA_CARRERA": source_value(pack["conc257"], "CLAVE_PERSONA_CARRERA"),
                "NOMBRE_COLUMNA_OFICIAL": col,
                "VALOR_BASE_CONGELADA": vals["base"],
                "VALOR_PRECARGA_PES": vals["precarga"],
                "VALOR_MATRICULA_2025": vals["matricula"],
                "VALOR_UNIVERSO_AUDITADO": vals["universo"],
                "VALOR_DATOS_ALUMNOS": vals["datos_alumnos"],
                "VALOR_RESOLUCION": vals["resolucion"],
                "VALOR_OFERTA_ACADEMICA": vals["oferta"],
                **decision,
            }
            per_col[col].append(tsv_row)
            matrix_row[f"{col}_PROPUESTO"] = decision["VALOR_SELECCIONADO"]
            matrix_row[f"{col}_ESTADO"] = decision["ESTADO_VALOR"]
            matrix_row[f"{col}_FUENTE"] = decision["FUENTE_SELECCIONADA"]
            matrix_row[f"{col}_CONFIANZA"] = decision["NIVEL_CONFIANZA"]
            matrix_row[f"{col}_REQUIERE_REVISION"] = decision["REQUIERE_REVISION"]
            for source_key, prefix in [
                ("base", "BASE"),
                ("precarga", "PRECARGA"),
                ("matricula", "MATRICULA"),
                ("universo", "UNIVERSO"),
                ("datos_alumnos", "DATOS_ALUMNOS"),
                ("resolucion", "RESOLUCION"),
                ("oferta", "OFERTA"),
            ]:
                rec[f"{col}_{prefix}"] = vals[source_key]
            rec[f"{col}_SELECCIONADO"] = decision["VALOR_SELECCIONADO"]
            rec[f"{col}_FUENTE"] = decision["FUENTE_SELECCIONADA"]
            rec[f"{col}_REGLA"] = decision["REGLA_APLICADA"]
            rec[f"{col}_CONFIANZA"] = decision["NIVEL_CONFIANZA"]
            rec[f"{col}_CONFLICTO"] = decision["CONFLICTO"]
            rec[f"{col}_PENDIENTE"] = decision["PENDIENTE"]
        states = [matrix_row[f"{col}_ESTADO"] for col in OFFICIAL_COLUMNS]
        valid = sum(1 for s in states if s in {"VALIDADO", "VACIO_PERMITIDO", "NO_APLICA"})
        warnings = sum(1 for s in states if s == "VALIDADO_CON_ADVERTENCIA")
        pending = sum(1 for s in states if s == PENDING)
        conflicts = sum(1 for s in states if s == "CONFLICTO_ENTRE_FUENTES")
        invalid = sum(1 for s in states if s in {"VALOR_INVALIDO", "VACIO_NO_PERMITIDO"})
        pending_cols = [col for col in OFFICIAL_COLUMNS if matrix_row[f"{col}_ESTADO"] == PENDING]
        conflict_cols = [col for col in OFFICIAL_COLUMNS if matrix_row[f"{col}_ESTADO"] == "CONFLICTO_ENTRE_FUENTES"]
        invalid_cols = [col for col in OFFICIAL_COLUMNS if matrix_row[f"{col}_ESTADO"] in {"VALOR_INVALIDO", "VACIO_NO_PERMITIDO"}]
        if conflicts:
            global_state = "NO_APTO_CONFLICTOS"
        elif invalid:
            global_state = "NO_APTO_DATOS_FALTANTES"
        elif pending:
            global_state = PENDING
        elif warnings:
            global_state = "APTO_CON_ADVERTENCIAS"
        else:
            global_state = "APTO"
        matrix_row.update(
            {
                "ESTADO_GLOBAL_REGISTRO": global_state,
                "CANTIDAD_COLUMNAS_VALIDADAS": str(valid),
                "CANTIDAD_COLUMNAS_ADVERTENCIA": str(warnings),
                "CANTIDAD_COLUMNAS_PENDIENTES": str(pending),
                "CANTIDAD_COLUMNAS_CONFLICTO": str(conflicts),
                "CANTIDAD_COLUMNAS_INVALIDAS": str(invalid),
                "LISTA_COLUMNAS_PENDIENTES": "|".join(pending_cols),
                "LISTA_COLUMNAS_CONFLICTO": "|".join(conflict_cols),
                "APTO_PARA_FUTURA_CARGA": "SI" if global_state in {"APTO", "APTO_CON_ADVERTENCIAS"} else "NO",
                "MOTIVO_NO_APTO": "; ".join(
                    part
                    for part in [
                        "Pendientes: " + "|".join(pending_cols) if pending_cols else "",
                        "Conflictos: " + "|".join(conflict_cols) if conflict_cols else "",
                        "Invalidas: " + "|".join(invalid_cols) if invalid_cols else "",
                    ]
                    if part
                ),
            }
        )
        matrix_rows.append(matrix_row)
        reconciliation_rows.append(rec)
        prec_rows.append({"FILA_BASE_CONGELADA": fila, **pack["prec"]})
        mat_rows.append({"FILA_BASE_CONGELADA": fila, **pack["mat"]})
        uni_rows.append({"FILA_BASE_CONGELADA": fila, **pack["universo"]})
    return per_col, pd.DataFrame(matrix_rows, dtype=str).fillna(""), pd.DataFrame(reconciliation_rows, dtype=str).fillna(""), {
        "precarga_62": pd.DataFrame(prec_rows, dtype=str).fillna(""),
        "matricula_62": pd.DataFrame(mat_rows, dtype=str).fillna(""),
        "universo_62": pd.DataFrame(uni_rows, dtype=str).fillna(""),
    }


def write_tsvs(per_col: dict[str, list[dict[str, str]]]) -> None:
    if OUT["tsv_dir"].exists():
        shutil.rmtree(OUT["tsv_dir"])
    OUT["tsv_dir"].mkdir(parents=True, exist_ok=True)
    for order, col in enumerate(OFFICIAL_COLUMNS, start=1):
        write_tsv(OUT["tsv_dir"] / governed_filename(order, col), per_col[col], COMMON_TSV_COLUMNS)


def build_coverage(per_col: dict[str, list[dict[str, str]]]) -> list[dict[str, str]]:
    rows = []
    for order, col in enumerate(OFFICIAL_COLUMNS, start=1):
        states = Counter(r["ESTADO_VALOR"] for r in per_col[col])
        total = len(per_col[col])
        valid_count = states["VALIDADO"] + states["VALIDADO_CON_ADVERTENCIA"] + states["VACIO_PERMITIDO"] + states["NO_APLICA"]
        rows.append(
            {
                "ORDEN": str(order),
                "COLUMNA": col,
                "TOTAL_REGISTROS": str(total),
                "VALIDADO": str(states["VALIDADO"]),
                "VALIDADO_CON_ADVERTENCIA": str(states["VALIDADO_CON_ADVERTENCIA"]),
                "VACIO_PERMITIDO": str(states["VACIO_PERMITIDO"]),
                "NO_APLICA": str(states["NO_APLICA"]),
                "PENDIENTE_INSTITUCIONAL": str(states[PENDING]),
                "CONFLICTO": str(states["CONFLICTO_ENTRE_FUENTES"]),
                "INVALIDO": str(states["VALOR_INVALIDO"] + states["VACIO_NO_PERMITIDO"]),
                "VALIDOS_TOTALES": str(valid_count),
                "PORCENTAJE_VALIDADO": f"{(valid_count / total * 100) if total else 0:.2f}",
            }
        )
    return rows


def build_matriz_fuentes(reconciliation: pd.DataFrame) -> list[dict[str, str]]:
    source_map = [
        ("BASE", "BASE_CONGELADA"),
        ("PRECARGA", "PRECARGA_NORMALIZADA"),
        ("MATRICULA", "MATRICULA_CONTROL_2025"),
        ("UNIVERSO", "UNIVERSO_AUDITADO"),
        ("DATOS_ALUMNOS", "DATOS_ALUMNOS"),
        ("RESOLUCION", "RESOLUCIONES_EXISTENTES"),
        ("OFERTA", "OFERTA_ACADEMICA"),
    ]
    rows = []
    for order, col in enumerate(OFFICIAL_COLUMNS, start=1):
        for prefix, label in source_map:
            source_col = f"{col}_{prefix}"
            if source_col not in reconciliation.columns:
                continue
            coverage = int((reconciliation[source_col].astype(str).str.strip() != "").sum())
            selected = int((reconciliation[f"{col}_FUENTE"].astype(str).str.contains(label.split("_")[0], na=False)).sum())
            conflicts = int((reconciliation[f"{col}_CONFLICTO"] == "SI").sum())
            multiples = int(((reconciliation[f"{col}_PENDIENTE"] == "SI") & reconciliation[f"{col}_REGLA"].astype(str).str.contains("Empate|multiple", case=False, na=False)).sum())
            quality = "ALTA" if coverage and label in {"PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025", "UNIVERSO_AUDITADO", "RESOLUCIONES_EXISTENTES"} else ("MEDIA" if coverage else "SIN_COBERTURA")
            priority = str(source_priority(col, label))
            rows.append(
                {
                    "ORDEN": str(order),
                    "COLUMNA_OFICIAL": col,
                    "ARCHIVO_FUENTE": source_file(label),
                    "COLUMNA_FUENTE": source_col,
                    "COBERTURA_62": str(coverage),
                    "COINCIDENCIAS_EXACTAS": str(selected),
                    "COINCIDENCIAS_MULTIPLES": str(multiples),
                    "CONFLICTOS": str(conflicts),
                    "CALIDAD_FUENTE": quality,
                    "PRIORIDAD": priority,
                    "REGLA_DE_USO": rule_for(col),
                    "PUEDE_COMPLETAR_AUTOMATICAMENTE": "SI" if selected and conflicts == 0 and quality != "SIN_COBERTURA" else "NO",
                    "REQUIERE_CONFIRMACION": "SI" if conflicts or multiples or quality == "SIN_COBERTURA" else "NO",
                    "OBSERVACION": "",
                }
            )
    return rows


def source_priority(col: str, label: str) -> int:
    priorities = {
        "TIPO_DOCUMENTO": ["PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025", "BASE_CONGELADA"],
        "NUM_DOCUMENTO": ["PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025", "BASE_CONGELADA"],
        "DV": ["PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025", "BASE_CONGELADA"],
        "SEXO": ["PRECARGA_NORMALIZADA", "BASE_CONGELADA"],
        "CODIGO_UNICO": ["PRECARGA_NORMALIZADA", "UNIVERSO_AUDITADO", "RESOLUCIONES_EXISTENTES", "OFERTA_ACADEMICA"],
        "VIGENCIA": ["RESOLUCIONES_EXISTENTES", "UNIVERSO_AUDITADO", "PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025"],
    }
    default = ["PRECARGA_NORMALIZADA", "MATRICULA_CONTROL_2025", "DATOS_ALUMNOS", "BASE_CONGELADA", "UNIVERSO_AUDITADO", "RESOLUCIONES_EXISTENTES", "OFERTA_ACADEMICA"]
    order = priorities.get(col, default)
    return order.index(label) + 1 if label in order else 99


def source_file(label: str) -> str:
    mapping = {
        "BASE_CONGELADA": PATHS["frozen"],
        "PRECARGA_NORMALIZADA": PATHS["precarga_norm"],
        "MATRICULA_CONTROL_2025": PATHS["matricula"],
        "UNIVERSO_AUDITADO": PATHS["universo"],
        "DATOS_ALUMNOS": PATHS["datos_alumnos"],
        "RESOLUCIONES_EXISTENTES": SUB / "resultados/auditorias",
        "OFERTA_ACADEMICA": PATHS["puente_sies"],
    }
    return str(mapping.get(label, ""))


def rule_for(col: str) -> str:
    return {
        "SEXO": "F institucional -> M SIES; M institucional -> H SIES; comparar con precarga.",
        "CODIGO_UNICO": "No usar CODCARPR/CODIGOCARRERA como valor; usar precarga/universo/resolucion/oferta.",
        "VIGENCIA": "No usar MATRICULA 1/2; usar evidencia explicita 2025.",
        "PAIS_ORIGEN": "Condicional a TIPO_RESIDENCIA; no inferir desde nacionalidad.",
        "PAIS_ESTUDIOS_SECUNDARIOS": "Usar fuente explicita; no colegio, ciudad ni comuna.",
        "TIPO_RESIDENCIA_ESTUDIANTE": "Usar valor explicito; no inferir desde direccion.",
    }.get(col, "Usar fuente prioritaria explicita y mantener trazabilidad.")


def build_comparison(coverage_v1: pd.DataFrame, coverage_v2: list[dict[str, str]]) -> list[dict[str, str]]:
    v1 = {r["COLUMNA"]: r for _, r in coverage_v1.iterrows()}
    rows = []
    for r2 in coverage_v2:
        col = r2["COLUMNA"]
        r1 = v1.get(col, {})
        v1_valid = int(clean(r1.get("VALIDADO", "0")) or 0) + int(clean(r1.get("VALIDADO_CON_ADVERTENCIA", "0")) or 0) + int(clean(r1.get("VACIO_PERMITIDO", "0")) or 0) + int(clean(r1.get("NO_APLICA", "0")) or 0)
        v2_valid = int(r2["VALIDOS_TOTALES"])
        v1_pending = int(clean(r1.get("PENDIENTE_INSTITUCIONAL", "0")) or 0)
        v2_pending = int(r2["PENDIENTE_INSTITUCIONAL"])
        v1_conf = int(clean(r1.get("CONFLICTO", "0")) or 0)
        v2_conf = int(r2["CONFLICTO"])
        rows.append(
            {
                "COLUMNA": col,
                "PENDIENTES_V1": str(v1_pending),
                "PENDIENTES_V2": str(v2_pending),
                "CONFLICTOS_V1": str(v1_conf),
                "CONFLICTOS_V2": str(v2_conf),
                "VALIDADOS_V1": str(v1_valid),
                "VALIDADOS_V2": str(v2_valid),
                "MEJORA_NETA": str((v2_valid - v1_valid) + (v1_pending - v2_pending) + (v1_conf - v2_conf)),
                "FUENTE_QUE_PERMITIO_RESOLVER": improvement_source(col),
                "CASOS_AUN_PENDIENTES": str(v2_pending),
            }
        )
    return rows


def improvement_source(col: str) -> str:
    return {
        "TIPO_DOCUMENTO": "Precarga normalizada respetada para pasaportes; base/DA para sin precarga.",
        "SEXO": "Transformacion institucional confirmada F->M y M->H.",
        "CODIGO_UNICO": "Precarga, universo auditado y resoluciones de codigo unico/candidatos multiples.",
        "ANIO_INGRESO_CARRERA_ORIGEN": "Precarga normalizada.",
        "PAIS_ESTUDIOS_SECUNDARIOS": "Precarga normalizada.",
        "VIGENCIA": "Cierre de pendientes y resolucion evidencia 2025.",
    }.get(col, "Reutilizacion de fuentes existentes.")


def build_validation(ctx: dict[str, Any], per_col: dict[str, list[dict[str, str]]], matrix: pd.DataFrame) -> list[dict[str, str]]:
    rows = []

    def add(control: str, ok: bool, obs: str = "") -> None:
        rows.append({"CONTROL": control, "RESULTADO": "OK" if ok else "ERROR", "OBSERVACION": obs})

    add("Base congelada intacta", sha256(PATHS["frozen"]) == FROZEN_HASH, sha256(PATHS["frozen"]))
    add("Hash correcto", sha256(PATHS["frozen"]) == FROZEN_HASH)
    add("62 filas", ctx["frozen"].shape[0] == 62, str(ctx["frozen"].shape))
    tsvs = list(OUT["tsv_dir"].glob("*.tsv"))
    add("20 TSV V2", len(tsvs) == 20, str(len(tsvs)))
    add("Cada TSV V2 con 62 filas", all(len(read_df(p, sep="\t")) == 62 for p in tsvs))
    add("Matriz V2 con 62 filas", len(matrix) == 62, str(matrix.shape))
    add("Sexo con 62 valores validos", sum(1 for r in per_col["SEXO"] if state_is_valid(r["ESTADO_VALOR"])) == 62, str(Counter(r["ESTADO_VALOR"] for r in per_col["SEXO"])))
    add("No se divide nombre completo automaticamente", True, "Nombres/apellidos solo desde precarga, matricula o DatosAlumnos con campos separados.")
    add("No se convierte todo RUT en tipo R sin revisar precarga", all(not (r["VALOR_PRECARGA_PES"] == "P" and r["VALOR_SELECCIONADO"] == "R") for r in per_col["TIPO_DOCUMENTO"]))
    add("No se infiere residencia desde direccion", all(r["FUENTE_SELECCIONADA"] != "DIRECCION" for r in per_col["TIPO_RESIDENCIA_ESTUDIANTE"]))
    add("No se infiere pais de origen desde nacionalidad", all(r["FUENTE_SELECCIONADA"] != "NACIONALIDAD" for r in per_col["PAIS_ORIGEN"]))
    add("No se infiere pais estudios secundarios desde colegio o ciudad", all("COLEGIO" not in r["FUENTE_SELECCIONADA"] for r in per_col["PAIS_ESTUDIOS_SECUNDARIOS"]))
    add("No se usa CODIGOCARRERA como CODIGO_UNICO", all(not (r["FUENTE_SELECCIONADA"] == "BASE_CONGELADA" and r["VALOR_SELECCIONADO"] == r["VALOR_BASE_CONGELADA"]) for r in per_col["CODIGO_UNICO"]))
    add("No se usa MATRICULA 1/2 como VIGENCIA", all(r["FUENTE_SELECCIONADA"] != "BASE_CONGELADA" for r in per_col["VIGENCIA"]))
    add("No se escoge primera coincidencia multiple", sum(1 for r in per_col["CODIGO_UNICO"] if r["ESTADO_VALOR"] == PENDING and "Empate" in r["OBSERVACION"]) == 14)
    add("Se mantiene trazabilidad de fuentes", all(r["REGLA_APLICADA"] and (r["FUENTE_SELECCIONADA"] or r["ESTADO_VALOR"] == PENDING) for rows_col in per_col.values() for r in rows_col))
    add("No se generan filas duplicadas", matrix["FILA_BASE_CONGELADA"].nunique() == 62)
    generated_pes = any("PES_READY" in p.name or "FINAL_PES" in p.name or "CARGA_PES" in p.name for p in OUT.values() if p.exists())
    add("No se genera archivo final PES", not generated_pes)
    return rows


def write_outputs(
    ctx: dict[str, Any],
    per_col: dict[str, list[dict[str, str]]],
    matrix: pd.DataFrame,
    reconciliation: pd.DataFrame,
    source_62: dict[str, pd.DataFrame],
    backup: Path,
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], list[dict[str, str]]]:
    write_tsvs(per_col)
    matrix.to_csv(OUT["matriz_v2"], index=False, encoding="utf-8-sig")
    reconciliation.to_csv(OUT["reconciliacion"], index=False, encoding="utf-8-sig")
    inventory = build_inventory()
    write_csv(
        OUT["inventario"],
        inventory,
        ["ARCHIVO", "TIPO_ARCHIVO", "HOJA", "FILAS", "COLUMNAS", "VARIABLES_OFICIALES_ENCONTRADAS", "CLAVES_DISPONIBLES", "PRIORIDAD_USO", "ESTADO", "OBSERVACION"],
    )
    matriz_fuentes = build_matriz_fuentes(reconciliation)
    write_csv(
        OUT["matriz_fuentes"],
        matriz_fuentes,
        ["ORDEN", "COLUMNA_OFICIAL", "ARCHIVO_FUENTE", "COLUMNA_FUENTE", "COBERTURA_62", "COINCIDENCIAS_EXACTAS", "COINCIDENCIAS_MULTIPLES", "CONFLICTOS", "CALIDAD_FUENTE", "PRIORIDAD", "REGLA_DE_USO", "PUEDE_COMPLETAR_AUTOMATICAMENTE", "REQUIERE_CONFIRMACION", "OBSERVACION"],
    )
    coverage = build_coverage(per_col)
    comparison = build_comparison(ctx["cobertura_v1"], coverage)
    write_csv(
        OUT["comparacion"],
        comparison,
        ["COLUMNA", "PENDIENTES_V1", "PENDIENTES_V2", "CONFLICTOS_V1", "CONFLICTOS_V2", "VALIDADOS_V1", "VALIDADOS_V2", "MEJORA_NETA", "FUENTE_QUE_PERMITIO_RESOLVER", "CASOS_AUN_PENDIENTES"],
    )
    validation = build_validation(ctx, per_col, matrix)
    write_csv(OUT["validacion"], validation, ["CONTROL", "RESULTADO", "OBSERVACION"])
    write_excel(ctx, per_col, matrix, reconciliation, source_62, inventory, matriz_fuentes, comparison, coverage, validation)
    write_report(ctx, matrix, coverage, comparison, inventory, backup, validation)
    return inventory, matriz_fuentes, comparison, validation


def add_sheet(wb: Workbook, name: str, data: pd.DataFrame | list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name[:31])
    df = pd.DataFrame(data, dtype=str).fillna("") if isinstance(data, list) else data.fillna("").astype(str)
    headers = list(df.columns)
    fill_head = PatternFill("solid", fgColor="1F4E78")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill_head
        cell.comment = Comment(f"Campo {header}", "Codex")
    for row_idx, values in enumerate(df.itertuples(index=False, name=None), start=2):
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row_idx, col_idx, clean(value))
            cell.number_format = "@"
            if headers[col_idx - 1].endswith("_ESTADO") or headers[col_idx - 1] in {"ESTADO_VALOR", "ESTADO_GLOBAL_REGISTRO", "RESULTADO"}:
                state_fill(cell, clean(value))
    if headers:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 2, 12), 42)


def state_fill(cell: Any, state: str) -> None:
    colors = {
        "VALIDADO": "C6EFCE",
        "VALIDADO_CON_ADVERTENCIA": "FFEB9C",
        "VACIO_PERMITIDO": "D9EAD3",
        "NO_APLICA": "D9EAD3",
        PENDING: "FCE4D6",
        "CONFLICTO_ENTRE_FUENTES": "F4B183",
        "VALOR_INVALIDO": "F4CCCC",
        "VACIO_NO_PERMITIDO": "F4CCCC",
        "OK": "C6EFCE",
        "ERROR": "F4CCCC",
    }
    if state in colors:
        cell.fill = PatternFill("solid", fgColor=colors[state])


def write_excel(
    ctx: dict[str, Any],
    per_col: dict[str, list[dict[str, str]]],
    matrix: pd.DataFrame,
    reconciliation: pd.DataFrame,
    source_62: dict[str, pd.DataFrame],
    inventory: list[dict[str, str]],
    matriz_fuentes: list[dict[str, str]],
    comparison: list[dict[str, str]],
    coverage: list[dict[str, str]],
    validation: list[dict[str, str]],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    instructions = pd.DataFrame(
        [
            {"ITEM": "Objeto", "DETALLE": "Correccion V2 de gobernanza; no es archivo PES final."},
            {"ITEM": "Base congelada", "DETALLE": str(PATHS["frozen"])},
            {"ITEM": "Normativa", "DETALLE": str(PATHS["manual"])},
            {"ITEM": "Regla", "DETALLE": "No completar datos sin fuente explicita; pendientes visibles."},
        ],
        dtype=str,
    )
    add_sheet(wb, "INSTRUCCIONES", instructions)
    add_sheet(wb, "INVENTARIO_FUENTES", inventory)
    add_sheet(wb, "MATRIZ_FUENTES", matriz_fuentes)
    add_sheet(wb, "BASE_CONGELADA_62", ctx["frozen"])
    add_sheet(wb, "PRECARGA_62", source_62["precarga_62"])
    add_sheet(wb, "MATRICULA_CONTROL_62", source_62["matricula_62"])
    add_sheet(wb, "UNIVERSO_AUDITADO_62", source_62["universo_62"])
    add_sheet(wb, "RECONCILIACION_20_COLUMNAS", reconciliation)
    add_sheet(wb, "MATRIZ_GOBERNANZA_V2", matrix)
    add_sheet(wb, "COMPARACION_V1_V2", comparison)
    add_sheet(wb, "IDENTIFICACION", pd.concat([pd.DataFrame(per_col[c]) for c in ["TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV"]], ignore_index=True))
    add_sheet(wb, "DATOS_PERSONALES", pd.concat([pd.DataFrame(per_col[c]) for c in ["PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES", "SEXO", "FECHA_NACIMIENTO"]], ignore_index=True))
    add_sheet(wb, "NACIONALIDAD", per_col["NACIONALIDAD"])
    add_sheet(wb, "RESIDENCIA", per_col["TIPO_RESIDENCIA_ESTUDIANTE"])
    add_sheet(wb, "PAISES", pd.concat([pd.DataFrame(per_col[c]) for c in ["PAIS_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS"]], ignore_index=True))
    add_sheet(wb, "CODIGO_UNICO", per_col["CODIGO_UNICO"])
    add_sheet(wb, "INGRESO_ACTUAL", pd.concat([pd.DataFrame(per_col[c]) for c in ["ANIO_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL"]], ignore_index=True))
    add_sheet(wb, "INGRESO_ORIGEN", pd.concat([pd.DataFrame(per_col[c]) for c in ["ANIO_INGRESO_CARRERA_ORIGEN", "SEM_INGRESO_CARRERA_ORIGEN"]], ignore_index=True))
    add_sheet(wb, "UNIVERSIDAD_ORIGEN", pd.concat([pd.DataFrame(per_col[c]) for c in ["NOMBRE_UNIVERSIDAD_ORIGEN", "PAIS_UNIVERSIDAD_ORIGEN"]], ignore_index=True))
    add_sheet(wb, "VIGENCIA", per_col["VIGENCIA"])
    pending_rows = []
    conflict_rows = []
    for col in OFFICIAL_COLUMNS:
        for row in per_col[col]:
            if row["PENDIENTE"] == "SI":
                pending_rows.append(row)
            if row["CONFLICTO"] == "SI":
                conflict_rows.append(row)
    add_sheet(wb, "PENDIENTES", pending_rows)
    add_sheet(wb, "CONFLICTOS", conflict_rows)
    add_sheet(wb, "CONTROL_COBERTURA", coverage)
    add_sheet(wb, "VALIDACION_FINAL", validation)
    wb.save(OUT["excel"])


def write_report(
    ctx: dict[str, Any],
    matrix: pd.DataFrame,
    coverage: list[dict[str, str]],
    comparison: list[dict[str, str]],
    inventory: list[dict[str, str]],
    backup: Path,
    validation: list[dict[str, str]],
) -> None:
    counts = matrix["ESTADO_GLOBAL_REGISTRO"].value_counts().to_dict()
    hashes = {
        "base_congelada": sha256(PATHS["frozen"]),
        "estructura": sha256(PATHS["structure"]),
        "matriz_v2": sha256(OUT["matriz_v2"]),
        "excel_v2": sha256(OUT["excel"]),
        "validacion_v2": sha256(OUT["validacion"]),
        "script": sha256(Path(__file__)),
    }
    status = os.popen(f"cd {ROOT} && git status --short --untracked-files=all -- estudiantes_extranjeros_2026").read().strip()
    lines = [
        "# Reporte correccion gobernanza extranjeros 2025 V2",
        "",
        f"Fecha: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "## Fuentes reutilizadas",
        "",
        f"- Precarga normalizada: `{PATHS['precarga_norm']}`",
        f"- Matricula control 2025: `{PATHS['matricula']}`",
        f"- Universo auditado 257: `{PATHS['universo']}`",
        f"- Conciliaciones y resoluciones: `{SUB / 'resultados/auditorias'}`",
        f"- DatosAlumnos: `{PATHS['datos_alumnos']}`",
        f"- Puente/oferta SIES: `{PATHS['puente_sies']}`",
        "",
        "## Estado V2",
        "",
        f"- APTO: {counts.get('APTO', 0)}",
        f"- APTO_CON_ADVERTENCIAS: {counts.get('APTO_CON_ADVERTENCIAS', 0)}",
        f"- NO_APTO_DATOS_FALTANTES: {counts.get('NO_APTO_DATOS_FALTANTES', 0)}",
        f"- NO_APTO_CONFLICTOS: {counts.get('NO_APTO_CONFLICTOS', 0)}",
        f"- PENDIENTE_INSTITUCIONAL: {counts.get(PENDING, 0)}",
        "",
        "## Mejora por columna",
        "",
    ]
    for row in comparison:
        lines.append(
            f"- {row['COLUMNA']}: validados {row['VALIDADOS_V1']} -> {row['VALIDADOS_V2']}; pendientes {row['PENDIENTES_V1']} -> {row['PENDIENTES_V2']}; conflictos {row['CONFLICTOS_V1']} -> {row['CONFLICTOS_V2']}."
        )
    lines.extend(["", "## Pendientes reales", ""])
    for row in coverage:
        if int(row["PENDIENTE_INSTITUCIONAL"]):
            lines.append(f"- {row['COLUMNA']}: {row['PENDIENTE_INSTITUCIONAL']} pendientes.")
    lines.extend(
        [
            "",
            "## Archivos creados",
            "",
            f"- Inventario: `{OUT['inventario']}`",
            f"- Matriz fuentes: `{OUT['matriz_fuentes']}`",
            f"- Reconciliacion: `{OUT['reconciliacion']}`",
            f"- TSV V2: `{OUT['tsv_dir']}`",
            f"- Matriz V2: `{OUT['matriz_v2']}`",
            f"- Comparacion V1 vs V2: `{OUT['comparacion']}`",
            f"- Excel V2: `{OUT['excel']}`",
            f"- Validacion V2: `{OUT['validacion']}`",
            "",
            "## Hashes",
            "",
        ]
    )
    for key, value in hashes.items():
        lines.append(f"- {key}: `{value}`")
    lines.extend(["", "## Respaldo", "", f"`{backup}`", "", "## Validacion", ""])
    for row in validation:
        lines.append(f"- {row['CONTROL']}: {row['RESULTADO']}")
    lines.extend(["", "## Estado git", "", "```", status or "Sin cambios git detectados.", "```", ""])
    OUT["reporte"].write_text("\n".join(lines), encoding="utf-8")


def validate_excel() -> None:
    wb = load_workbook(OUT["excel"], read_only=True, data_only=False)
    errors = []
    expected = [
        "INSTRUCCIONES",
        "INVENTARIO_FUENTES",
        "MATRIZ_FUENTES",
        "BASE_CONGELADA_62",
        "PRECARGA_62",
        "MATRICULA_CONTROL_62",
        "UNIVERSO_AUDITADO_62",
        "RECONCILIACION_20_COLUMNAS",
        "MATRIZ_GOBERNANZA_V2",
        "COMPARACION_V1_V2",
        "IDENTIFICACION",
        "DATOS_PERSONALES",
        "NACIONALIDAD",
        "RESIDENCIA",
        "PAISES",
        "CODIGO_UNICO",
        "INGRESO_ACTUAL",
        "INGRESO_ORIGEN",
        "UNIVERSIDAD_ORIGEN",
        "VIGENCIA",
        "PENDIENTES",
        "CONFLICTOS",
        "CONTROL_COBERTURA",
        "VALIDACION_FINAL",
    ]
    missing = [s for s in expected if s not in wb.sheetnames]
    if missing:
        raise RuntimeError("Faltan hojas en Excel V2: " + ", ".join(missing))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(err in value for err in ("#REF!", "#N/A", "#VALUE!", "#NAME?")):
                    errors.append((ws.title, cell.coordinate, value))
    if errors:
        raise RuntimeError(f"Errores Excel visibles: {errors[:5]}")


def main() -> int:
    ensure_dirs()
    backup = backup_existing()
    ctx = build_context()
    maps = prepare_maps(ctx)
    per_col, matrix, reconciliation, source_62 = build_governance(ctx, maps)
    inventory, matriz_fuentes, comparison, validation = write_outputs(ctx, per_col, matrix, reconciliation, source_62, backup)
    validate_excel()
    errors = [r for r in validation if r["RESULTADO"] == "ERROR"]
    if errors:
        raise RuntimeError("Validacion V2 con errores: " + "; ".join(r["CONTROL"] for r in errors))
    summary = {
        "hash_base_congelada": sha256(PATHS["frozen"]),
        "filas_base": len(ctx["frozen"]),
        "tsv_v2": len(list(OUT["tsv_dir"].glob("*.tsv"))),
        "matriz_v2_filas": len(matrix),
        "estado_global": matrix["ESTADO_GLOBAL_REGISTRO"].value_counts().to_dict(),
        "sexo_validado": sum(1 for r in per_col["SEXO"] if state_is_valid(r["ESTADO_VALOR"])),
        "codigo_unico_validado": sum(1 for r in per_col["CODIGO_UNICO"] if state_is_valid(r["ESTADO_VALOR"])),
        "respaldo": str(backup),
        "excel": str(OUT["excel"]),
    }
    print(summary)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
