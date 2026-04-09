#!/usr/bin/env python3
"""Genera la base maestra auditable de extranjeros regulares 2025.

No genera CSV final para PES. Usa el diagnostico inicial como antecedente y
materializa una base de trabajo, auditorias, planilla de gestion y manifiesto.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[2]
SUB = ROOT / "estudiantes_extranjeros_2026"
DOCS = SUB / "docs"
DATA_INTERIM = SUB / "data" / "interim"
DATA_PROCESSED = SUB / "data" / "processed"
AUDITORIAS = SUB / "resultados" / "auditorias"
REPORTES = SUB / "resultados" / "reportes"
GESTION = SUB / "resultados" / "archivos_gestion"
BACKUPS = SUB / "backups"

RUN_TS = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
RUN_STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")

OFFICIAL_FIELDS = [
    "TIPO_DOCUMENTO",
    "NUM_DOCUMENTO",
    "DV",
    "PRIMER_APELLIDO",
    "SEGUNDO_APELLIDO",
    "NOMBRES",
    "SEXO",
    "FECHA_NACIMIENTO",
    "NACIONALIDAD",
    "TIPO_RESIDENCIA_ESTUDIANTE",
    "PAIS_DE_ORIGEN",
    "PAIS_ESTUDIOS_SECUNDARIOS",
    "CODIGO_UNICO",
    "ANIO_INGRESO_CARRERA_ACTUAL",
    "SEM_INGRESO_CARRERA_ACTUAL",
    "ANIO_INGRESO_CARRERA_ORIGEN",
    "SEM_INGRESO_CARRERA_ORIGEN",
    "NOMBRE_UNIVERSIDAD_ORIGEN",
    "PAIS_UNIVERSIDAD_ORIGEN",
    "VIGENCIA",
]

TRACE_SUFFIXES = [
    "BRUTO",
    "NORMALIZADO",
    "FUENTE_CAMPO",
    "ARCHIVO_ORIGEN",
    "HOJA_ORIGEN",
    "COLUMNA_ORIGEN",
    "REGLA_APLICADA",
    "ESTADO_CAMPO",
    "OBSERVACION_CAMPO",
]

CONTROL_COLUMNS = [
    "ID_CONTROL",
    "CODIGO_IES_NUM",
    "CODCLI",
    "CLAVE_PERSONA_INTERNA",
    "CLAVE_REGISTRO",
    "CLASIFICACION_UNIVERSO",
    "ESTADO_REGISTRO",
    "REQUIERE_GESTION_MANUAL",
    "CANTIDAD_CAMPOS_FALTANTES",
    "CANTIDAD_CONFLICTOS",
    "FUENTE_PRIORITARIA",
    "FECHA_ACTUALIZACION",
]

EXTRA_COLUMNS = [
    "NOMBRE_COMPLETO",
    "NOMBRE_CARRERA",
    "CODCARPR",
    "SEDE",
    "JORNADA",
    "MODALIDAD",
    "NIVEL",
    "FECHA_MATRICULA",
    "ANOMATRICULA",
    "PERIODOMATRICULA",
    "ESTADO_ACADEMICO",
    "SITUACION",
    "MATRICULA",
    "FUENTE_VIGENCIA",
    "NACIONALIDAD_BRUTA",
    "PAIS_ESTUDIOS_SECUNDARIOS_CANDIDATO",
    "CODIGO_UNICO_ESTADO_RESOLUCION",
    "CODIGO_UNICO_CODIGOS_CANDIDATOS",
    "OBSERVACION_GENERAL",
]


def ensure_dirs() -> None:
    for path in [DATA_INTERIM, DATA_PROCESSED, AUDITORIAS, REPORTES, GESTION, BACKUPS]:
        path.mkdir(parents=True, exist_ok=True)


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT))
    except Exception:
        return str(path)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text


def upper_ascii(value: Any) -> str:
    text = clean(value).upper()
    text = text.translate(str.maketrans("ÁÉÍÓÚÜÑ", "AEIOUUN"))
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_doc(value: Any) -> str:
    return re.sub(r"[.\-,\s]", "", clean(value).upper())


def split_run(value: Any) -> tuple[str, str]:
    text = normalize_doc(value)
    if len(text) >= 2 and text[-1] in "0123456789K" and text[:-1].isdigit():
        return text[:-1], text[-1]
    return text, ""


def normalize_date(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    for dayfirst in (True, False):
        parsed = pd.to_datetime(text, errors="coerce", dayfirst=dayfirst)
        if not pd.isna(parsed):
            return parsed.strftime("%d-%m-%Y")
    return text


def normalize_sex(value: Any) -> tuple[str, str]:
    raw = clean(value).upper()
    if raw == "F":
        return "M", "Sexo institucional F transformado a M (Mujer) segun catalogo SIES."
    if raw == "M":
        return "H", "Sexo institucional M transformado a H (Hombre) segun catalogo SIES."
    if raw in {"H", "NB"}:
        return raw, "Valor ya en catalogo SIES."
    return "", "Valor de sexo no mapeado con seguridad."


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: "" if row.get(k) is None else row.get(k) for k in fieldnames})


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def backup_targets(paths: list[Path]) -> Path:
    backup_dir = BACKUPS / f"pre_fase2_{RUN_STAMP}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for path in paths:
        if not path.exists():
            continue
        target = backup_dir / path.name
        shutil.copy2(path, target)
        rows.append(
            {
                "archivo_original": rel(path),
                "archivo_respaldo": rel(target),
                "sha256_original": sha256(path),
                "tamano_bytes": path.stat().st_size,
                "fecha_respaldo": RUN_TS,
            }
        )
    manifest = backup_dir / "MANIFIESTO_RESPALDO_PRE_FASE2.csv"
    write_csv(manifest, rows, ["archivo_original", "archivo_respaldo", "sha256_original", "tamano_bytes", "fecha_respaldo"])
    return backup_dir


def read_inputs() -> dict[str, Any]:
    da_path = ROOT / "input" / "PROMEDIOSDEALUMNOS_7804.xlsx"
    da = pd.read_excel(da_path, sheet_name="DatosAlumnos", dtype=str, engine="openpyxl")
    da.columns = [clean(c) for c in da.columns]

    base_datos = pd.read_excel(da_path, sheet_name="base_datos", dtype=str, engine="openpyxl")
    base_datos.columns = [clean(c) for c in base_datos.columns]

    matriz = pd.read_excel(da_path, sheet_name="matriz", dtype=str, engine="openpyxl")
    matriz.columns = [clean(c) for c in matriz.columns]

    gob = pd.read_csv(ROOT / "gobernanza_nac.tsv", sep="\t", dtype=str)
    gob["NACIONALIDAD_JOIN"] = gob["NACIONALIDAD_ORIG"].map(upper_ascii)

    bridge = pd.read_csv(ROOT / "control" / "catalogos" / "PUENTE_SIES_COMPILADO.tsv", sep="\t", dtype=str)

    ac = pd.read_csv(ROOT / "resultados" / "matricula_avance_curricular_2025_control.csv", dtype=str)
    ac["NUM_DOCUMENTO_NORM"] = ac["NUM_DOCUMENTO"].map(normalize_doc)

    dur_path = ROOT / "DURACION_ESTUDIOS.tsv"
    duracion = pd.read_csv(dur_path, sep="\t", dtype=str) if dur_path.exists() else pd.DataFrame()

    normalized_ok = read_archivo_listo()

    conflictos_path = AUDITORIAS / "CONFLICTOS_IDENTIFICADORES_Y_CRUCES.csv"
    conflictos = pd.read_csv(conflictos_path, dtype=str) if conflictos_path.exists() else pd.DataFrame()

    return {
        "da": da,
        "base_datos": base_datos,
        "matriz": matriz,
        "gob": gob,
        "bridge": bridge,
        "ac": ac,
        "duracion": duracion,
        "normalized_ok": normalized_ok,
        "conflictos": conflictos,
    }


def read_archivo_listo() -> pd.DataFrame:
    path = ROOT / "resultados" / "archivo_listo_para_sies.xlsx"
    if not path.exists():
        return pd.DataFrame()
    cols = [
        "TIPO_DOC",
        "N_DOC",
        "DV",
        "PRIMER_APELLIDO",
        "SEGUNDO_APELLIDO",
        "NOMBRE",
        "SEXO",
        "FECH_NAC",
        "NAC",
        "PAIS_EST_SEC",
        "CODIGO_CARRERA_SIES_FINAL",
        "ANIO_ING_ACT",
        "SEM_ING_ACT",
        "ANIO_ING_ORI",
        "SEM_ING_ORI",
        "VIG",
        "CODCLI",
        "DA_ANOMATRICULA",
        "INCLUIR_EN_MATRICULA_32",
        "ESTADO_CARGA_PREGRADO",
        "NAC_STATUS",
        "PAIS_EST_SEC_STATUS",
    ]
    df = pd.read_excel(path, sheet_name="ARCHIVO_LISTO_SUBIDA", dtype=str, usecols=lambda c: c in cols, engine="openpyxl")
    df = df[
        df.get("DA_ANOMATRICULA", "").astype(str).str.strip().eq("2025")
        & df.get("NAC", "").astype(str).str.strip().ne("38")
        & df.get("NAC", "").notna()
        & df.get("INCLUIR_EN_MATRICULA_32", "").astype(str).str.upper().eq("SI")
    ].copy()
    if df.empty:
        return df
    df = df.sort_values(["CODCLI"]).drop_duplicates("CODCLI", keep="first")
    return df


def build_country_catalog() -> list[dict[str, Any]]:
    text = (DOCS / "Instructivo_Estudiantes_Extranjeros_SIES_2026.txt").read_text(encoding="utf-8", errors="replace")
    section = text.split("ANEXO III: Tabla países", 1)[-1].split("ANEXO IV:", 1)[0]
    names: dict[int, str] = {}
    for line in section.splitlines():
        line = upper_ascii(line)
        if not re.search(r"\d", line):
            continue
        # Captura pares codigo/nombre dentro de filas de 4 columnas.
        matches = list(re.finditer(r"(?<!\d)(\d{1,3})\s+([A-Z][A-Z\s\-\.'/]+?)(?=\s+\d{1,3}\s+[A-Z]|$)", line))
        for m in matches:
            code = int(m.group(1))
            if 1 <= code <= 197:
                name = clean(m.group(2))
                name = re.sub(r"\s+", " ", name).strip(" -")
                if name and not name.startswith("PAGINA"):
                    names.setdefault(code, name)
    supplements = {
        4: "ANDORRA",
        6: "ANTIGUA Y BARBUDA",
        24: "BOSNIA-HERZEGOVINA",
        44: "COREA DEL NORTE",
        48: "CROACIA",
        55: "EMIRATOS ARABES UNIDOS",
        148: "REPUBLICA DEMOCRATICA DEL CONGO",
        156: "SAN VICENTE Y LAS GRANADINAS",
    }
    names.update(supplements)
    rows = []
    for code in range(1, 198):
        name = names.get(code, f"PAIS CODIGO {code}")
        rows.append(
            {
                "CODIGO_PAIS": code,
                "NOMBRE_PAIS": name,
                "ETIQUETA": f"{code} - {name}",
                "PERMITIDO_NACIONALIDAD": "NO" if code == 38 else "SI",
                "PERMITIDO_PAIS_ORIGEN": "NO" if code == 38 else "SI",
                "FUENTE": "Instructivo Anexo III; nombre completado cuando el texto extraido lo permite",
            }
        )
    return rows


def classify_candidates(inputs: dict[str, Any]) -> pd.DataFrame:
    da = inputs["da"].copy()
    da["NACIONALIDAD_NORM"] = da["NACIONALIDAD"].map(upper_ascii)
    base = da[
        da["ANOMATRICULA"].astype(str).str.strip().eq("2025")
        & da["NACIONALIDAD_NORM"].ne("")
        & da["NACIONALIDAD_NORM"].ne("NAN")
        & ~da["NACIONALIDAD_NORM"].str.contains("CHIL", na=False)
    ].copy()
    base["RUN_CUERPO"] = base["RUT"].map(lambda x: split_run(x)[0])
    base["DV_NORM"] = base["RUT"].map(lambda x: split_run(x)[1])
    base["TIPO_DOCUMENTO_CANDIDATO"] = base["RUN_CUERPO"].map(lambda x: "R" if x else "")
    base["NOMBRE_COMPLETO"] = (
        base.get("NOMBRES", "").fillna("").astype(str).str.strip()
        + " "
        + base.get("APELLIDO PATERNO", "").fillna("").astype(str).str.strip()
        + " "
        + base.get("APELLIDO MATERNO", "").fillna("").astype(str).str.strip()
    ).str.replace(r"\s+", " ", regex=True).str.strip()

    gob = inputs["gob"]
    base = base.merge(
        gob[["NACIONALIDAD_JOIN", "COD_NAC", "PAIS_NORM", "ESTADO_GOBERNANZA"]],
        how="left",
        left_on="NACIONALIDAD_NORM",
        right_on="NACIONALIDAD_JOIN",
    )

    mobility_re = re.compile(r"INTERCAMBIO|MOVILIDAD|PASANT|STUDY ABROAD|CONVENIO INTERNACIONAL|PROGRAMA CORTO")
    classifications = []
    motives = []
    for _, row in base.iterrows():
        blob = " ".join(clean(row.get(c, "")) for c in ["NOMBRE_L", "CATEGORIA", "VIASDEADMISION", "MOTIVODESELECCION", "NOMBREUNIVERSIDAD", "CARRERARANTERIOR"]).upper()
        if mobility_re.search(blob):
            classifications.append("POSIBLE_INTERCAMBIO")
            motives.append("Indicio textual explicito de movilidad/intercambio/pasantia/convenio; requiere revision.")
        elif clean(row.get("COD_NAC", "")) and clean(row.get("COD_NAC", "")) != "38":
            classifications.append("EXTRANJERO_CONFIRMADO")
            motives.append("Nacionalidad extranjera mapeada a codigo oficial distinto de Chile.")
        elif upper_ascii(row.get("NACIONALIDAD", "")) in {"POR DEFINIR", "SIN INFORMACION", "NO INFORMADO"} or not clean(row.get("COD_NAC", "")):
            classifications.append("NACIONALIDAD_POR_CONFIRMAR")
            motives.append("Nacionalidad no mapeada o ambigua; no se confirma como extranjero hasta respuesta institucional.")
        else:
            classifications.append("NO_CORRESPONDE")
            motives.append("Fuente revisada no acredita nacionalidad extranjera.")
    base["CLASIFICACION_UNIVERSO"] = classifications
    base["MOTIVO_CLASIFICACION"] = motives
    base = base.sort_values(["CODCLI", "RUN_CUERPO"]).reset_index(drop=True)
    base["ID_CONTROL"] = [f"EE2025-{i:04d}" for i in range(1, len(base) + 1)]
    return base


def make_universe_rows(candidates: pd.DataFrame) -> list[dict[str, Any]]:
    rows = []
    for _, row in candidates.iterrows():
        rows.append(
            {
                "ID_CONTROL": row["ID_CONTROL"],
                "CODCLI": row.get("CODCLI", ""),
                "TIPO_DOCUMENTO_BRUTO": "RUT" if clean(row.get("RUT", "")) else "",
                "NUM_DOCUMENTO_BRUTO": row.get("RUT", ""),
                "DV_BRUTO": row.get("DV_NORM", ""),
                "NOMBRE_COMPLETO": row.get("NOMBRE_COMPLETO", ""),
                "NACIONALIDAD_BRUTA": row.get("NACIONALIDAD", ""),
                "NACIONALIDAD_CODIGO": row.get("COD_NAC", ""),
                "CLASIFICACION_UNIVERSO": row.get("CLASIFICACION_UNIVERSO", ""),
                "MOTIVO_CLASIFICACION": row.get("MOTIVO_CLASIFICACION", ""),
                "FUENTE_PRINCIPAL": "input/PROMEDIOSDEALUMNOS_7804.xlsx::DatosAlumnos",
                "FUENTES_COMPLEMENTARIAS": "gobernanza_nac.tsv; resultados/archivo_listo_para_sies.xlsx; control/catalogos/PUENTE_SIES_COMPILADO.tsv",
                "ESTADO_REVISION": "PRELIMINAR_AUDITABLE",
            }
        )
    return rows


def parse_codes(text: Any) -> list[str]:
    raw = clean(text)
    if not raw:
        return []
    return sorted({x.strip() for x in re.split(r"\s*\|\s*", raw) if x.strip()})


def resolve_codigo(row: pd.Series, inputs: dict[str, Any], norm_by_codcli: dict[str, pd.Series]) -> dict[str, Any]:
    codcli = clean(row.get("CODCLI", ""))
    codcarpr = clean(row.get("CODCARPR", ""))
    jornada = clean(row.get("JORNADA", ""))
    bridge = inputs["bridge"]
    prev_conf = inputs["conflictos"]
    conflict_types = []
    if not prev_conf.empty:
        hits = prev_conf[prev_conf["CLAVE"].astype(str).str.strip() == codcli]
        conflict_types = sorted(set(hits["TIPO_CONFLICTO"].dropna().astype(str)))

    norm_row = norm_by_codcli.get(codcli)
    if norm_row is not None and clean(norm_row.get("CODIGO_CARRERA_SIES_FINAL", "")):
        code = clean(norm_row.get("CODIGO_CARRERA_SIES_FINAL", ""))
        return {
            "CODIGO_UNICO": code,
            "CLASIFICACION": "RESUELTO_REGLA_INSTITUCIONAL" if conflict_types else "RESUELTO_UNIVOCO",
            "FUENTE": "resultados/archivo_listo_para_sies.xlsx::ARCHIVO_LISTO_SUBIDA",
            "METODO": "CODCLI en salida MU auditada OK 2025",
            "CODIGOS_CANDIDATOS": code,
            "CONFLICTOS_PREVIOS": " | ".join(conflict_types),
            "OBSERVACION": "Codigo reutilizado desde fuente normalizada validada; no es CSV final PES.",
        }

    m = bridge[bridge["CODCARPR"].fillna("").map(clean) == codcarpr].copy()
    mj = m[m["JORNADA"].fillna("").map(clean) == jornada].copy()
    exact_codes = sorted({clean(x) for x in mj.get("CODIGO_UNICO_FINAL", pd.Series(dtype=str)) if clean(x)})
    exact_potential = sorted({clean(x) for x in mj.get("CODIGOS_SIES_POTENCIALES", pd.Series(dtype=str)) if clean(x)})
    all_codes = sorted({clean(x) for x in m.get("CODIGO_UNICO_FINAL", pd.Series(dtype=str)) if clean(x)})
    if len(exact_codes) == 1:
        return {
            "CODIGO_UNICO": exact_codes[0],
            "CLASIFICACION": "RESUELTO_REGLA_INSTITUCIONAL" if conflict_types else "RESUELTO_UNIVOCO",
            "FUENTE": "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
            "METODO": "CODCARPR + JORNADA con un CODIGO_UNICO_FINAL",
            "CODIGOS_CANDIDATOS": " | ".join(exact_codes),
            "CONFLICTOS_PREVIOS": " | ".join(conflict_types),
            "OBSERVACION": "Seleccion estructural; no se eligio por similitud de nombre.",
        }
    if exact_potential:
        return {
            "CODIGO_UNICO": "",
            "CLASIFICACION": "AMBIGUO",
            "FUENTE": "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
            "METODO": "CODCARPR + JORNADA con multiples potenciales o sin final",
            "CODIGOS_CANDIDATOS": " | ".join(exact_potential),
            "CONFLICTOS_PREVIOS": " | ".join(conflict_types),
            "OBSERVACION": "Requiere confirmacion; no hay regla deterministica segura.",
        }
    if all_codes:
        return {
            "CODIGO_UNICO": "",
            "CLASIFICACION": "REQUIERE_CONFIRMACION",
            "FUENTE": "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
            "METODO": "CODCARPR sin coincidencia unica por jornada",
            "CODIGOS_CANDIDATOS": " | ".join(all_codes),
            "CONFLICTOS_PREVIOS": " | ".join(conflict_types),
            "OBSERVACION": "Existe codigo para CODCARPR, pero no coincide de forma segura con jornada.",
        }
    return {
        "CODIGO_UNICO": "",
        "CLASIFICACION": "SIN_CODIGO",
        "FUENTE": "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
        "METODO": "Sin match estructural en puente",
        "CODIGOS_CANDIDATOS": "",
        "CONFLICTOS_PREVIOS": " | ".join(conflict_types),
        "OBSERVACION": "Evaluar codigo temporal si la oferta no existe en SIES.",
    }


def set_field(base: dict[str, Any], field: str, raw: Any, norm: Any, source: str, file: str, sheet: str, column: str, rule: str, status: str, obs: str) -> None:
    base[f"{field}_BRUTO"] = clean(raw)
    base[f"{field}_NORMALIZADO"] = clean(norm)
    base[f"{field}_FUENTE_CAMPO"] = source
    base[f"{field}_ARCHIVO_ORIGEN"] = file
    base[f"{field}_HOJA_ORIGEN"] = sheet
    base[f"{field}_COLUMNA_ORIGEN"] = column
    base[f"{field}_REGLA_APLICADA"] = rule
    base[f"{field}_ESTADO_CAMPO"] = status
    base[f"{field}_OBSERVACION_CAMPO"] = obs
    base[field] = clean(norm)


def build_base(candidates: pd.DataFrame, inputs: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    norm_by_codcli = {clean(r["CODCLI"]): r for _, r in inputs["normalized_ok"].iterrows()} if not inputs["normalized_ok"].empty else {}
    ac = inputs["ac"]
    ac_codes = ac.groupby("NUM_DOCUMENTO_NORM")["CODIGO_UNICO"].apply(lambda s: sorted({clean(x) for x in s if clean(x)})).to_dict()
    cod_ies = "162"
    if not inputs["matriz"].empty and "CODIGO_IES_NUM" in inputs["matriz"]:
        vals = [clean(x) for x in inputs["matriz"]["CODIGO_IES_NUM"] if clean(x)]
        if vals:
            cod_ies = vals[0]

    base_rows: list[dict[str, Any]] = []
    resolucion_rows: list[dict[str, Any]] = []
    vigencia_rows: list[dict[str, Any]] = []

    for _, src in candidates.iterrows():
        row: dict[str, Any] = {c: "" for c in CONTROL_COLUMNS + EXTRA_COLUMNS + OFFICIAL_FIELDS}
        for field in OFFICIAL_FIELDS:
            for suffix in TRACE_SUFFIXES:
                row[f"{field}_{suffix}"] = ""

        id_control = clean(src["ID_CONTROL"])
        codcli = clean(src.get("CODCLI", ""))
        row["ID_CONTROL"] = id_control
        row["CODIGO_IES_NUM"] = cod_ies
        row["CODCLI"] = codcli
        row["CLAVE_PERSONA_INTERNA"] = f"CODCLI:{codcli}|DOC:R:{src.get('RUN_CUERPO','')}"
        row["CLASIFICACION_UNIVERSO"] = src.get("CLASIFICACION_UNIVERSO", "")
        row["FUENTE_PRIORITARIA"] = "input/PROMEDIOSDEALUMNOS_7804.xlsx::DatosAlumnos"
        row["FECHA_ACTUALIZACION"] = RUN_TS
        row["NOMBRE_COMPLETO"] = src.get("NOMBRE_COMPLETO", "")
        row["NOMBRE_CARRERA"] = src.get("NOMBRE_L", "")
        row["CODCARPR"] = src.get("CODCARPR", "")
        row["SEDE"] = src.get("SEDE", "")
        row["JORNADA"] = src.get("JORNADA", "")
        row["MODALIDAD"] = src.get("MODALIDAD", "")
        row["NIVEL"] = src.get("NIVEL", "")
        row["FECHA_MATRICULA"] = normalize_date(src.get("FECHAMATRICULA", ""))
        row["ANOMATRICULA"] = src.get("ANOMATRICULA", "")
        row["PERIODOMATRICULA"] = src.get("PERIODOMATRICULA", "")
        row["ESTADO_ACADEMICO"] = src.get("ESTADOACADEMICO", "")
        row["SITUACION"] = src.get("SITUACION", "")
        row["MATRICULA"] = src.get("MATRICULA", "")
        row["NACIONALIDAD_BRUTA"] = src.get("NACIONALIDAD", "")

        set_field(row, "TIPO_DOCUMENTO", "RUT" if src.get("RUT", "") else "", "R" if src.get("RUN_CUERPO", "") else "", "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "RUT", "RUN presente se informa como R; IPE no permitido.", "COMPLETO_TRANSFORMADO", "")
        set_field(row, "NUM_DOCUMENTO", src.get("RUT", ""), src.get("RUN_CUERPO", ""), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "RUT", "Eliminar puntuacion y separar DV.", "COMPLETO_TRANSFORMADO", "CODCLI no reemplaza NUM_DOCUMENTO.")
        set_field(row, "DV", src.get("RUT", ""), src.get("DV_NORM", ""), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "RUT", "Separar digito verificador del RUN.", "COMPLETO_TRANSFORMADO", "")
        set_field(row, "PRIMER_APELLIDO", src.get("APELLIDO PATERNO", ""), upper_ascii(src.get("APELLIDO PATERNO", "")), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "APELLIDO PATERNO", "Mayusculas, sin acentos, trim y espacios simples.", "COMPLETO_TRANSFORMADO", "")
        set_field(row, "SEGUNDO_APELLIDO", src.get("APELLIDO MATERNO", ""), upper_ascii(src.get("APELLIDO MATERNO", "")), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "APELLIDO MATERNO", "Mayusculas, sin acentos, trim y espacios simples.", "COMPLETO_TRANSFORMADO", "")
        set_field(row, "NOMBRES", src.get("NOMBRES", ""), upper_ascii(src.get("NOMBRES", "")), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "NOMBRES", "Mayusculas, sin acentos, trim y espacios simples.", "COMPLETO_TRANSFORMADO", "")
        sexo_norm, sexo_obs = normalize_sex(src.get("SEXO", ""))
        sexo_status = "COMPLETO_TRANSFORMADO" if sexo_norm else "PENDIENTE_GESTION"
        set_field(row, "SEXO", src.get("SEXO", ""), sexo_norm, "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "SEXO", "Mapeo a catalogo SIES M/H/NB.", sexo_status, sexo_obs)
        set_field(row, "FECHA_NACIMIENTO", src.get("FECHANACIMIENTO", ""), normalize_date(src.get("FECHANACIMIENTO", "")), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "FECHANACIMIENTO", "Normalizar fecha a DD-MM-AAAA.", "COMPLETO_TRANSFORMADO", "")

        nac_status = "COMPLETO_VALIDADO" if clean(src.get("COD_NAC", "")) and clean(src.get("COD_NAC", "")) != "38" else "PENDIENTE_GESTION"
        set_field(row, "NACIONALIDAD", src.get("NACIONALIDAD", ""), src.get("COD_NAC", ""), "DatosAlumnos + gobernanza_nac.tsv", "gobernanza_nac.tsv", "", "NACIONALIDAD_ORIG/COD_NAC", "Mapeo gobernado a codigo 1..197 distinto de Chile.", nac_status, src.get("ESTADO_GOBERNANZA", "No mapeado"))

        set_field(row, "TIPO_RESIDENCIA_ESTUDIANTE", "", "", "No detectada", "", "", "", "No inferir desde domicilio, nacionalidad o modalidad.", "PENDIENTE_GESTION", "Campo critico requerido por instructivo.")
        set_field(row, "PAIS_DE_ORIGEN", "", "", "No detectada", "", "", "", "Depende de TIPO_RESIDENCIA_ESTUDIANTE; no inferir.", "PENDIENTE_GESTION", "Debe quedar vacio si residencia=1; requerido si residencia=2 o 3.")

        norm_row = norm_by_codcli.get(codcli)
        if norm_row is not None and clean(norm_row.get("PAIS_EST_SEC", "")):
            set_field(row, "PAIS_ESTUDIOS_SECUNDARIOS", norm_row.get("PAIS_EST_SEC", ""), norm_row.get("PAIS_EST_SEC", ""), "archivo_listo_para_sies normalizado", "resultados/archivo_listo_para_sies.xlsx", "ARCHIVO_LISTO_SUBIDA", "PAIS_EST_SEC", "Reutilizar normalizacion MU ya incluida en registro OK 2025.", "COMPLETO_VALIDADO", clean(norm_row.get("PAIS_EST_SEC_STATUS", "")))
            row["PAIS_ESTUDIOS_SECUNDARIOS_CANDIDATO"] = norm_row.get("PAIS_EST_SEC", "")
        else:
            set_field(row, "PAIS_ESTUDIOS_SECUNDARIOS", f"{src.get('COMUNACOLEGIO','')} / {src.get('CIUDADCOLEGIO','')}", "", "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "COMUNACOLEGIO/CIUDADCOLEGIO", "No completar automaticamente; requiere pais donde completo y aprobo secundaria.", "PENDIENTE_GESTION", "Dato de colegio/localidad queda solo como antecedente.")

        code_info = resolve_codigo(src, inputs, norm_by_codcli)
        code_status = "COMPLETO_VALIDADO" if code_info["CODIGO_UNICO"] else "CONFLICTO_REQUIERE_CONFIRMACION"
        if code_info["CLASIFICACION"] in {"SIN_CODIGO"}:
            code_status = "PENDIENTE_GESTION"
        set_field(row, "CODIGO_UNICO", code_info["CODIGOS_CANDIDATOS"], code_info["CODIGO_UNICO"], code_info["FUENTE"], code_info["FUENTE"].split("::")[0], code_info["FUENTE"].split("::")[1] if "::" in code_info["FUENTE"] else "", "CODIGO_UNICO_FINAL/CODIGO_CARRERA_SIES_FINAL", code_info["METODO"], code_status, code_info["OBSERVACION"])
        row["CODIGO_UNICO_ESTADO_RESOLUCION"] = code_info["CLASIFICACION"]
        row["CODIGO_UNICO_CODIGOS_CANDIDATOS"] = code_info["CODIGOS_CANDIDATOS"]

        set_field(row, "ANIO_INGRESO_CARRERA_ACTUAL", src.get("ANOINGRESO", ""), src.get("ANOINGRESO", ""), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "ANOINGRESO", "Copia directa; validar rango antes de carga.", "COMPLETO_VALIDADO", "")
        set_field(row, "SEM_INGRESO_CARRERA_ACTUAL", src.get("PERIODOINGRESO", ""), src.get("PERIODOINGRESO", ""), "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "PERIODOINGRESO", "Copia directa; validar 1/2.", "COMPLETO_VALIDADO", "")

        if norm_row is not None and clean(norm_row.get("ANIO_ING_ORI", "")):
            set_field(row, "ANIO_INGRESO_CARRERA_ORIGEN", norm_row.get("ANIO_ING_ORI", ""), norm_row.get("ANIO_ING_ORI", ""), "archivo_listo_para_sies normalizado", "resultados/archivo_listo_para_sies.xlsx", "ARCHIVO_LISTO_SUBIDA", "ANIO_ING_ORI", "Reutilizar motor MU normalizado.", "COMPLETO_VALIDADO", "")
            set_field(row, "SEM_INGRESO_CARRERA_ORIGEN", norm_row.get("SEM_ING_ORI", ""), norm_row.get("SEM_ING_ORI", ""), "archivo_listo_para_sies normalizado", "resultados/archivo_listo_para_sies.xlsx", "ARCHIVO_LISTO_SUBIDA", "SEM_ING_ORI", "Reutilizar motor MU normalizado.", "COMPLETO_VALIDADO", "")
        else:
            set_field(row, "ANIO_INGRESO_CARRERA_ORIGEN", "", "", "No resuelto con seguridad", "", "", "", "No copiar anio actual sin regla de ingreso/origen trazada.", "PENDIENTE_GESTION", "")
            set_field(row, "SEM_INGRESO_CARRERA_ORIGEN", "", "", "No resuelto con seguridad", "", "", "", "No copiar semestre actual sin regla de ingreso/origen trazada.", "PENDIENTE_GESTION", "")

        set_field(row, "NOMBRE_UNIVERSIDAD_ORIGEN", src.get("NOMBREUNIVERSIDAD", ""), "", "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "NOMBREUNIVERSIDAD", "Campo solo aplica a doble titulacion; no se completa automaticamente.", "NO_APLICA_PENDIENTE_CONFIRMACION", "La fuente puede representar institucion anterior, no doble titulacion.")
        set_field(row, "PAIS_UNIVERSIDAD_ORIGEN", "", "", "No detectada", "", "", "", "Campo solo aplica a doble titulacion.", "NO_APLICA_PENDIENTE_CONFIRMACION", "")

        vigencia = "1" if clean(src.get("ANOMATRICULA", "")) == "2025" else ""
        vig_status = "COMPLETO_VALIDADO" if vigencia else "PENDIENTE_CONFIRMACION"
        evidencia = f"ANOMATRICULA={src.get('ANOMATRICULA','')}; PERIODOMATRICULA={src.get('PERIODOMATRICULA','')}; FECHAMATRICULA={normalize_date(src.get('FECHAMATRICULA',''))}; ESTADOACADEMICO={src.get('ESTADOACADEMICO','')}; MATRICULA={src.get('MATRICULA','')}"
        set_field(row, "VIGENCIA", evidencia, vigencia, "DatosAlumnos", "input/PROMEDIOSDEALUMNOS_7804.xlsx", "DatosAlumnos", "ANOMATRICULA/PERIODOMATRICULA/FECHAMATRICULA", "Mantener solo con evidencia de matricula 2025; no usar vigencia 2026.", vig_status, "Vigencia propuesta refiere al proceso extranjeros 2025.")
        row["FUENTE_VIGENCIA"] = evidencia

        row["CLAVE_REGISTRO"] = f"{cod_ies}|{row['TIPO_DOCUMENTO']}|{row['NUM_DOCUMENTO']}|{row['CODIGO_UNICO'] or 'SIN_CODIGO'}"

        conflict_count = 0
        missing_count = 0
        for field in OFFICIAL_FIELDS:
            status = row.get(f"{field}_ESTADO_CAMPO", "")
            if status.startswith("PENDIENTE"):
                missing_count += 1
            if "CONFLICTO" in status:
                conflict_count += 1
        # Pais de origen se mantiene como gestion pendiente por condicion de residencia desconocida.
        row["CANTIDAD_CAMPOS_FALTANTES"] = missing_count
        row["CANTIDAD_CONFLICTOS"] = conflict_count
        row["REQUIERE_GESTION_MANUAL"] = "SI" if missing_count or conflict_count or row["CLASIFICACION_UNIVERSO"] != "EXTRANJERO_CONFIRMADO" else "NO"
        row["ESTADO_REGISTRO"] = "PENDIENTE_GESTION_MANUAL" if row["REQUIERE_GESTION_MANUAL"] == "SI" else "LISTO_TRANSFORMACION_FUTURA"
        if row["CLASIFICACION_UNIVERSO"] == "POSIBLE_INTERCAMBIO":
            row["ESTADO_REGISTRO"] = "REVISAR_POSIBLE_INTERCAMBIO"
        row["OBSERVACION_GENERAL"] = src.get("MOTIVO_CLASIFICACION", "")

        base_rows.append(row)

        resolucion_rows.append(make_resolution_row(src, row, inputs, code_info, ac_codes))
        vigencia_rows.append(
            {
                "CODCLI": codcli,
                "NUM_DOCUMENTO": row["NUM_DOCUMENTO"],
                "CODIGO_UNICO": row["CODIGO_UNICO"],
                "VIGENCIA_PROPUESTA": row["VIGENCIA"] if row["VIGENCIA"] else "PENDIENTE_CONFIRMACION",
                "EVIDENCIA_2025": evidencia,
                "ARCHIVO_FUENTE": "input/PROMEDIOSDEALUMNOS_7804.xlsx::DatosAlumnos",
                "ESTADO_VALIDACION": "EVIDENCIA_MATRICULA_2025" if row["VIGENCIA"] == "1" else "SIN_EVIDENCIA",
                "REQUIERE_CONFIRMACION": "SI" if row["CLASIFICACION_UNIVERSO"] != "EXTRANJERO_CONFIRMADO" else "NO",
                "OBSERVACION": "No se uso vigencia 2026.",
            }
        )

    return base_rows, resolucion_rows, vigencia_rows


def make_resolution_row(src: pd.Series, base_row: dict[str, Any], inputs: dict[str, Any], code_info: dict[str, Any], ac_codes: dict[str, list[str]]) -> dict[str, Any]:
    dur = inputs["duracion"]
    code = base_row["CODIGO_UNICO"]
    oferta = ""
    if code and not dur.empty and "CODIGO_UNICO" in dur:
        hit = dur[dur["CODIGO_UNICO"].fillna("").map(clean) == code]
        if not hit.empty:
            h = hit.iloc[0]
            oferta = f"NOMBRE_CARRERA={clean(h.get('NOMBRE_CARRERA',''))}; JORNADA={clean(h.get('JORNADA',''))}; MODALIDAD={clean(h.get('MODALIDAD',''))}; VIGENCIA={clean(h.get('VIGENCIA',''))}"
    return {
        "ID_CONTROL": base_row["ID_CONTROL"],
        "CODCLI": base_row["CODCLI"],
        "TIPO_CONFLICTO_PREVIO": code_info["CONFLICTOS_PREVIOS"],
        "CODCARPR": src.get("CODCARPR", ""),
        "SEDE": src.get("SEDE", ""),
        "JORNADA": src.get("JORNADA", ""),
        "MODALIDAD": src.get("MODALIDAD", ""),
        "VERSION": code[-1:] if code else "",
        "NOMBRE_OFICIAL_PROGRAMA": src.get("NOMBRE_L", ""),
        "OFERTA_ACADEMICA_REFERENCIA": oferta,
        "CODIGO_UNICO_DISPONIBLE": code,
        "CODIGOS_CANDIDATOS": code_info["CODIGOS_CANDIDATOS"],
        "MATRICULA_EFECTIVA_2025": "SI" if clean(src.get("ANOMATRICULA", "")) == "2025" else "PENDIENTE",
        "CODIGOS_AVANCE_2025": " | ".join(ac_codes.get(clean(src.get("RUN_CUERPO", "")), [])),
        "CLASIFICACION_RESOLUCION": code_info["CLASIFICACION"],
        "FUENTE_RESOLUCION": code_info["FUENTE"],
        "METODO_RESOLUCION": code_info["METODO"],
        "OBSERVACION": code_info["OBSERVACION"],
    }


def make_cross_audit(candidates: pd.DataFrame, inputs: dict[str, Any], base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    docs_unique = candidates.groupby("RUN_CUERPO")["CODCLI"].nunique()
    rows.append(
        {
            "METODO": "DOCUMENTO_EXACTO_NORMALIZADO",
            "CLAVE_UTILIZADA": "TIPO_DOCUMENTO + NUM_DOCUMENTO + DV",
            "CANTIDAD_COINCIDENCIAS": len(candidates),
            "UNICIDAD": "UNICO" if int((docs_unique > 1).sum()) == 0 else "CONFLICTO",
            "CONFLICTO": int((docs_unique > 1).sum()),
            "FUENTE_GANADORA": "input/PROMEDIOSDEALUMNOS_7804.xlsx::DatosAlumnos",
            "FUENTES_DESCARTADAS": "",
            "MOTIVO_SELECCION": "Fuente principal con ANOMATRICULA=2025 y documento disponible.",
        }
    )
    bd = inputs["base_datos"].copy()
    bd["RUN_CUERPO"] = bd["RUT"].map(normalize_doc) if "RUT" in bd else ""
    matched_bd = candidates["RUN_CUERPO"].isin(set(bd["RUN_CUERPO"].dropna().astype(str))).sum() if "RUN_CUERPO" in bd else 0
    rows.append(
        {
            "METODO": "CODCLI_MEDIANTE_TABLA_MAESTRA",
            "CLAVE_UTILIZADA": "RUN/DV -> CODCLI en hoja base_datos",
            "CANTIDAD_COINCIDENCIAS": int(matched_bd),
            "UNICIDAD": "PARCIAL",
            "CONFLICTO": 0,
            "FUENTE_GANADORA": "input/PROMEDIOSDEALUMNOS_7804.xlsx::base_datos",
            "FUENTES_DESCARTADAS": "",
            "MOTIVO_SELECCION": "Tabla de equivalencias disponible pero parcial; CODCLI no reemplaza documento oficial.",
        }
    )
    norm = inputs["normalized_ok"]
    rows.append(
        {
            "METODO": "CODCLI_ARCHIVO_NORMALIZADO_OK",
            "CLAVE_UTILIZADA": "CODCLI",
            "CANTIDAD_COINCIDENCIAS": int(candidates["CODCLI"].isin(set(norm["CODCLI"])).sum()) if not norm.empty else 0,
            "UNICIDAD": "PARCIAL",
            "CONFLICTO": 0,
            "FUENTE_GANADORA": "resultados/archivo_listo_para_sies.xlsx::ARCHIVO_LISTO_SUBIDA",
            "FUENTES_DESCARTADAS": "",
            "MOTIVO_SELECCION": "Solo se reutilizan filas con INCLUIR_EN_MATRICULA_32=SI y ANOMATRICULA=2025.",
        }
    )
    resolved = sum(1 for r in base_rows if clean(r.get("CODIGO_UNICO", "")))
    unresolved = len(base_rows) - resolved
    rows.append(
        {
            "METODO": "CODCARPR_JORNADA_PUENTE_SIES",
            "CLAVE_UTILIZADA": "CODCARPR + JORNADA",
            "CANTIDAD_COINCIDENCIAS": resolved,
            "UNICIDAD": "MIXTA",
            "CONFLICTO": unresolved,
            "FUENTE_GANADORA": "control/catalogos/PUENTE_SIES_COMPILADO.tsv",
            "FUENTES_DESCARTADAS": "Coincidencias por nombre solamente",
            "MOTIVO_SELECCION": "Se acepta solo codigo final unico o fuente normalizada OK; ambiguos pasan a gestion.",
        }
    )
    name_key = candidates["NOMBRE_COMPLETO"].map(upper_ascii) + "|" + candidates["FECHANACIMIENTO"].map(normalize_date)
    dup_name = int(name_key.value_counts().gt(1).sum())
    rows.append(
        {
            "METODO": "NOMBRE_FECHA_SOLO_DIAGNOSTICO",
            "CLAVE_UTILIZADA": "NOMBRE_COMPLETO + FECHA_NACIMIENTO",
            "CANTIDAD_COINCIDENCIAS": len(candidates),
            "UNICIDAD": "UNICO" if dup_name == 0 else "REVISAR",
            "CONFLICTO": dup_name,
            "FUENTE_GANADORA": "No aplica",
            "FUENTES_DESCARTADAS": "No se usa para completar datos",
            "MOTIVO_SELECCION": "Metodo reservado para diagnostico; nunca completa automaticamente.",
        }
    )
    return rows


def coverage_rows(base_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    total = len(base_rows)
    rows = []
    for field in OFFICIAL_FIELDS:
        complete = 0
        pending = 0
        conflicts = 0
        no_aplica = 0
        source_counter = Counter()
        for row in base_rows:
            status = row.get(f"{field}_ESTADO_CAMPO", "")
            if "CONFLICTO" in status:
                conflicts += 1
            elif status.startswith("PENDIENTE"):
                pending += 1
            elif status.startswith("NO_APLICA"):
                no_aplica += 1
            elif clean(row.get(field, "")):
                complete += 1
            else:
                pending += 1
            source_counter[row.get(f"{field}_FUENTE_CAMPO", "")] += 1
        source = source_counter.most_common(1)[0][0] if source_counter else ""
        action = "Sin accion" if pending == 0 and conflicts == 0 else "Gestion manual / resolver conflicto"
        denom = total - no_aplica if total > no_aplica else total
        pct = (complete / denom * 100) if denom else 0
        rows.append(
            {
                "VARIABLE": field,
                "TOTAL_REGISTROS": total,
                "COMPLETOS": complete,
                "PENDIENTES": pending,
                "CONFLICTOS": conflicts,
                "NO_APLICA": no_aplica,
                "PORCENTAJE_COBERTURA": f"{pct:.1f}%",
                "FUENTE_PRINCIPAL": source,
                "ACCION_PENDIENTE": action,
            }
        )
    return rows


def manual_rows(base_rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    simple = []
    detail = []
    for row in base_rows:
        pending_fields = []
        for field in OFFICIAL_FIELDS:
            status = row.get(f"{field}_ESTADO_CAMPO", "")
            if status.startswith("PENDIENTE") or "CONFLICTO" in status:
                pending_fields.append(field)
                detail.append(
                    {
                        "ID_CONTROL": row["ID_CONTROL"],
                        "CODCLI": row["CODCLI"],
                        "NOMBRE_COMPLETO": row["NOMBRE_COMPLETO"],
                        "CARRERA": row["NOMBRE_CARRERA"],
                        "SEDE": row["SEDE"],
                        "CAMPO": field,
                        "ESTADO_CAMPO": status,
                        "MOTIVO": row.get(f"{field}_OBSERVACION_CAMPO", ""),
                        "FUENTE_REVISADA": row.get(f"{field}_FUENTE_CAMPO", ""),
                    }
                )
        if pending_fields:
            simple.append(
                {
                    "CODCLI": row["CODCLI"],
                    "NOMBRE_COMPLETO": row["NOMBRE_COMPLETO"],
                    "CARRERA": row["NOMBRE_CARRERA"],
                    "SEDE": row["SEDE"],
                    "RUN_O_PASAPORTE": row["NUM_DOCUMENTO"],
                    "NACIONALIDAD": row["NACIONALIDAD"],
                    "TIPO_RESIDENCIA_PENDIENTE": "SI" if "TIPO_RESIDENCIA_ESTUDIANTE" in pending_fields else "NO",
                    "PAIS_ORIGEN_PENDIENTE": "SI" if "PAIS_DE_ORIGEN" in pending_fields else "NO",
                    "PAIS_ESTUDIOS_SECUNDARIOS_PENDIENTE": "SI" if "PAIS_ESTUDIOS_SECUNDARIOS" in pending_fields else "NO",
                    "OTRO_CAMPO_PENDIENTE": " | ".join([f for f in pending_fields if f not in {"TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_DE_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS"}]),
                    "OBSERVACION": row["OBSERVACION_GENERAL"],
                }
            )
    return simple, detail


def create_management_workbook(path: Path, base_rows: list[dict[str, Any]], countries: list[dict[str, Any]], coverage: list[dict[str, Any]], resolution: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "INSTRUCCIONES"
    lines = [
        "Objetivo: completar datos pendientes para Estudiantes Extranjeros SIES 2026, matricula efectiva 2025.",
        "No usar IPE. Los tipos validos son R RUN y P Pasaporte.",
        "Debe informarse RUN o pasaporte. CODCLI es interno y no reemplaza NUM_DOCUMENTO.",
        "Nacionalidad, pais de origen y pais de estudios secundarios son conceptos distintos.",
        "No completar datos por supuestos.",
        "Pais de estudios secundarios: pais donde completo y aprobo la ensenanza media/secundaria.",
        "Tipo de residencia: situacion previa al inicio de estudios superiores en Chile.",
        "Si residencia es 1, PAIS_DE_ORIGEN debe quedar vacio.",
        "Si residencia es 2 o 3, PAIS_DE_ORIGEN es obligatorio y no puede ser Chile.",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("ESTUDIANTES")
    headers = [
        "ID_CONTROL",
        "CODCLI",
        "NOMBRE_COMPLETO",
        "CARRERA",
        "SEDE",
        "TIPO_DOCUMENTO_EXISTENTE",
        "NUM_DOCUMENTO_EXISTENTE",
        "NACIONALIDAD_EXISTENTE",
        "NACIONALIDAD_RESPUESTA",
        "TIPO_RESIDENCIA_ESTUDIANTE_RESPUESTA",
        "PAIS_DE_ORIGEN_RESPUESTA",
        "PAIS_ESTUDIOS_SECUNDARIOS_RESPUESTA",
        "CONFIRMA_NACIONALIDAD",
        "CONFIRMA_MATRICULA_2025",
        "OBSERVACIONES_DOCENCIA",
        "RESPONSABLE_RESPUESTA",
        "FECHA_RESPUESTA",
        "CAMPOS_PENDIENTES",
        "OBSERVACION_INTERNA",
    ]
    ws.append(headers)
    pending_rows = [r for r in base_rows if r["REQUIERE_GESTION_MANUAL"] == "SI"]
    for row in pending_rows:
        pending = []
        for field in OFFICIAL_FIELDS:
            status = row.get(f"{field}_ESTADO_CAMPO", "")
            if status.startswith("PENDIENTE") or "CONFLICTO" in status:
                pending.append(field)
        ws.append(
            [
                row["ID_CONTROL"],
                row["CODCLI"],
                row["NOMBRE_COMPLETO"],
                row["NOMBRE_CARRERA"],
                row["SEDE"],
                row["TIPO_DOCUMENTO"],
                row["NUM_DOCUMENTO"],
                row["NACIONALIDAD"],
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                " | ".join(pending),
                row["OBSERVACION_GENERAL"],
            ]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(wrap_text=True)
    widths = {
        "A": 16,
        "B": 18,
        "C": 34,
        "D": 42,
        "E": 12,
        "F": 18,
        "G": 20,
        "H": 18,
        "I": 26,
        "J": 34,
        "K": 32,
        "L": 36,
        "M": 20,
        "N": 22,
        "O": 36,
        "P": 24,
        "Q": 18,
        "R": 54,
        "S": 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    pending_fill = PatternFill("solid", fgColor="FFF2CC")
    conflict_fill = PatternFill("solid", fgColor="F4CCCC")
    valid_fill = PatternFill("solid", fgColor="D9EAD3")
    for row_idx in range(2, ws.max_row + 1):
        pending_text = str(ws.cell(row_idx, 18).value or "")
        if "NACIONALIDAD" in pending_text:
            ws.cell(row_idx, 9).fill = pending_fill
        else:
            ws.cell(row_idx, 9).fill = valid_fill
        if "TIPO_RESIDENCIA_ESTUDIANTE" in pending_text:
            ws.cell(row_idx, 10).fill = pending_fill
        if "PAIS_DE_ORIGEN" in pending_text:
            ws.cell(row_idx, 11).fill = pending_fill
        if "PAIS_ESTUDIOS_SECUNDARIOS" in pending_text:
            ws.cell(row_idx, 12).fill = pending_fill
        else:
            ws.cell(row_idx, 12).fill = valid_fill
        if "CODIGO_UNICO" in str(ws.cell(row_idx, 18).value or ""):
            for col_idx in range(1, ws.max_column + 1):
                if col_idx in {18, 19}:
                    ws.cell(row_idx, col_idx).fill = conflict_fill
    dv_res = DataValidation(type="list", formula1='"1 - Con residencia previa en Chile,2 - Sin residencia previa en Chile,3 - No reside en Chile"', allow_blank=True)
    dv_si_no = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
    ws.add_data_validation(dv_res)
    ws.add_data_validation(dv_si_no)
    dv_res.add(f"J2:J{max(ws.max_row, 2)}")
    dv_si_no.add(f"M2:N{max(ws.max_row, 2)}")

    ws_p = wb.create_sheet("LISTA_PAISES")
    country_headers = ["CODIGO_PAIS", "NOMBRE_PAIS", "ETIQUETA", "PERMITIDO_NACIONALIDAD", "PERMITIDO_PAIS_ORIGEN", "FUENTE"]
    ws_p.append(country_headers)
    for c in countries:
        ws_p.append([c[h] for h in country_headers])
    for cell in ws_p[1]:
        cell.font = Font(bold=True)
    ws_p.column_dimensions["A"].width = 14
    ws_p.column_dimensions["B"].width = 36
    ws_p.column_dimensions["C"].width = 44
    last_country = ws_p.max_row
    country_formula = f"=LISTA_PAISES!$C$2:$C${last_country}"
    dv_country = DataValidation(type="list", formula1=country_formula, allow_blank=True)
    ws.add_data_validation(dv_country)
    dv_country.add(f"I2:I{max(ws.max_row, 2)}")
    dv_country.add(f"K2:K{max(ws.max_row, 2)}")
    dv_country.add(f"L2:L{max(ws.max_row, 2)}")

    ws = wb.create_sheet("DEFINICIONES")
    defs = [
        ("TIPO_RESIDENCIA_ESTUDIANTE=1", "Extranjero con residencia previa en Chile."),
        ("TIPO_RESIDENCIA_ESTUDIANTE=2", "Extranjero sin residencia previa en Chile."),
        ("TIPO_RESIDENCIA_ESTUDIANTE=3", "Extranjero que no reside en Chile."),
        ("PAIS_DE_ORIGEN", "Pais donde residia habitualmente si no tenia residencia previa, o pais donde reside si no reside en Chile. No usar Chile."),
        ("PAIS_ESTUDIOS_SECUNDARIOS", "Pais donde completo y aprobo la ensenanza media/secundaria."),
    ]
    ws.append(["CONCEPTO", "DEFINICION"])
    for row in defs:
        ws.append(row)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 110

    ws = wb.create_sheet("CONTROL_COBERTURA")
    cov_headers = list(coverage[0].keys()) if coverage else []
    ws.append(cov_headers)
    for r in coverage:
        ws.append([r[h] for h in cov_headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws = wb.create_sheet("CONFLICTOS_CODIGO_UNICO")
    res_headers = list(resolution[0].keys()) if resolution else []
    ws.append(res_headers)
    for r in resolution:
        if r["CLASIFICACION_RESOLUCION"] not in {"RESUELTO_UNIVOCO", "RESUELTO_REGLA_INSTITUCIONAL"} or r["TIPO_CONFLICTO_PREVIO"]:
            ws.append([r[h] for h in res_headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws = wb.create_sheet("POSIBLES_INTERCAMBIO")
    pos = [r for r in base_rows if r["CLASIFICACION_UNIVERSO"] == "POSIBLE_INTERCAMBIO"]
    headers_pos = ["ID_CONTROL", "CODCLI", "NOMBRE_COMPLETO", "NOMBRE_CARRERA", "MOTIVO"]
    ws.append(headers_pos)
    for r in pos:
        ws.append([r["ID_CONTROL"], r["CODCLI"], r["NOMBRE_COMPLETO"], r["NOMBRE_CARRERA"], r["OBSERVACION_GENERAL"]])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws = wb["ESTUDIANTES"]
    # Comentario orientador en columnas de respuesta.
    ws["J1"].comment = Comment("Obligatorio para todos los registros pendientes.", "Codex")
    ws["K1"].comment = Comment("Debe quedar vacio si residencia=1; obligatorio y distinto de Chile si residencia=2 o 3.", "Codex")
    ws["L1"].comment = Comment("Solo completar cuando este pendiente; corresponde al pais donde completo y aprobo secundaria.", "Codex")

    wb.save(path)


def manifest(paths: list[Path], backup_dir: Path) -> list[dict[str, Any]]:
    rows = []
    for path in paths:
        if not path.exists():
            continue
        rows.append(
            {
                "RUTA_RELATIVA": rel(path),
                "NOMBRE": path.name,
                "TAMANO_BYTES": path.stat().st_size,
                "SHA256": sha256(path),
                "FECHA_GENERACION": RUN_TS,
                "ORIGEN": "fase2_base_maestra",
                "OBSERVACION": "No es CSV final PES",
            }
        )
    rows.append(
        {
            "RUTA_RELATIVA": rel(backup_dir),
            "NOMBRE": backup_dir.name,
            "TAMANO_BYTES": "",
            "SHA256": "",
            "FECHA_GENERACION": RUN_TS,
            "ORIGEN": "respaldo_pre_sobrescritura",
            "OBSERVACION": "Contiene copias solo si existian archivos de salida previos.",
        }
    )
    return rows


def report_text(base_rows: list[dict[str, Any]], universe_rows: list[dict[str, Any]], coverage: list[dict[str, Any]], resolution: list[dict[str, Any]], vigencia: list[dict[str, Any]], backup_dir: Path) -> str:
    cls = Counter(r["CLASIFICACION_UNIVERSO"] for r in universe_rows)
    res = Counter(r["CLASIFICACION_RESOLUCION"] for r in resolution)
    vig = Counter(r["VIGENCIA_PROPUESTA"] for r in vigencia)
    cov_lines = "\n".join(
        f"| {r['VARIABLE']} | {r['TOTAL_REGISTROS']} | {r['COMPLETOS']} | {r['PENDIENTES']} | {r['CONFLICTOS']} | {r['NO_APLICA']} | {r['PORCENTAJE_COBERTURA']} | {r['FUENTE_PRINCIPAL']} | {r['ACCION_PENDIENTE']} |"
        for r in coverage
    )
    return f"""# Reporte Fase 2 - Base Maestra Extranjeros Regulares 2025

Fecha de ejecucion: {RUN_TS}

## Universo preliminar

- Total candidatos: {len(universe_rows)}
- Extranjeros confirmados: {cls.get('EXTRANJERO_CONFIRMADO', 0)}
- Nacionalidad por confirmar: {cls.get('NACIONALIDAD_POR_CONFIRMAR', 0)}
- Posible intercambio: {cls.get('POSIBLE_INTERCAMBIO', 0)}
- No corresponde: {cls.get('NO_CORRESPONDE', 0)}

## Base maestra

- Registros en base maestra: {len(base_rows)}
- Registros que requieren gestion manual: {sum(1 for r in base_rows if r['REQUIERE_GESTION_MANUAL'] == 'SI')}
- Archivo: `data/interim/BASE_MAESTRA_EXTRANJEROS_REGULARES_2025.csv`

## Codigo unico

{json.dumps(res, ensure_ascii=False, indent=2)}

## Vigencia 2025

{json.dumps(vig, ensure_ascii=False, indent=2)}

La vigencia propuesta se basa en evidencia 2025 de DatosAlumnos; no se uso vigencia 2026.

## Respaldo

`{rel(backup_dir)}`

## Tabla de control por variable

| Variable | Total | Completos | Pendientes | Conflictos | No aplica | Cobertura | Fuente principal | Accion pendiente |
|---|---:|---:|---:|---:|---:|---:|---|---|
{cov_lines}
"""


def main() -> int:
    ensure_dirs()
    outputs = [
        AUDITORIAS / "UNIVERSO_CANDIDATOS_EXTRANJEROS_2025.csv",
        DATA_INTERIM / "BASE_MAESTRA_EXTRANJEROS_REGULARES_2025.csv",
        AUDITORIAS / "AUDITORIA_CRUCES_BASE_MAESTRA.csv",
        AUDITORIAS / "RESOLUCION_CODIGO_UNICO.csv",
        AUDITORIAS / "AUDITORIA_VIGENCIA_2025.csv",
        GESTION / "PLANILLA_GESTION_DOCENCIA_EXTRANJEROS_2025.xlsx",
        AUDITORIAS / "NOMINA_SOLICITUD_DATOS_DOCENCIA.csv",
        AUDITORIAS / "NOMINA_GESTION_MANUAL_DETALLE_FASE2.csv",
        AUDITORIAS / "COBERTURA_VARIABLES_BASE_MAESTRA.csv",
        AUDITORIAS / "CATALOGO_PAISES_SIES_2026.csv",
        REPORTES / "REPORTE_BASE_MAESTRA_FASE2.md",
        AUDITORIAS / "MANIFIESTO_ARCHIVOS_FASE2.csv",
    ]
    backup_dir = backup_targets(outputs)

    inputs = read_inputs()
    countries = build_country_catalog()
    candidates = classify_candidates(inputs)
    universe = make_universe_rows(candidates)
    base_rows, resolution, vigencia = build_base(candidates, inputs)
    cross_audit = make_cross_audit(candidates, inputs, base_rows)
    coverage = coverage_rows(base_rows)
    simple_manual, detail_manual = manual_rows(base_rows)

    trace_columns = [f"{field}_{suffix}" for field in OFFICIAL_FIELDS for suffix in TRACE_SUFFIXES]
    base_columns = CONTROL_COLUMNS + trace_columns + OFFICIAL_FIELDS + EXTRA_COLUMNS

    write_csv(AUDITORIAS / "UNIVERSO_CANDIDATOS_EXTRANJEROS_2025.csv", universe, ["ID_CONTROL", "CODCLI", "TIPO_DOCUMENTO_BRUTO", "NUM_DOCUMENTO_BRUTO", "DV_BRUTO", "NOMBRE_COMPLETO", "NACIONALIDAD_BRUTA", "NACIONALIDAD_CODIGO", "CLASIFICACION_UNIVERSO", "MOTIVO_CLASIFICACION", "FUENTE_PRINCIPAL", "FUENTES_COMPLEMENTARIAS", "ESTADO_REVISION"])
    write_csv(DATA_INTERIM / "BASE_MAESTRA_EXTRANJEROS_REGULARES_2025.csv", base_rows, base_columns)
    write_csv(AUDITORIAS / "AUDITORIA_CRUCES_BASE_MAESTRA.csv", cross_audit, ["METODO", "CLAVE_UTILIZADA", "CANTIDAD_COINCIDENCIAS", "UNICIDAD", "CONFLICTO", "FUENTE_GANADORA", "FUENTES_DESCARTADAS", "MOTIVO_SELECCION"])
    write_csv(AUDITORIAS / "RESOLUCION_CODIGO_UNICO.csv", resolution, ["ID_CONTROL", "CODCLI", "TIPO_CONFLICTO_PREVIO", "CODCARPR", "SEDE", "JORNADA", "MODALIDAD", "VERSION", "NOMBRE_OFICIAL_PROGRAMA", "OFERTA_ACADEMICA_REFERENCIA", "CODIGO_UNICO_DISPONIBLE", "CODIGOS_CANDIDATOS", "MATRICULA_EFECTIVA_2025", "CODIGOS_AVANCE_2025", "CLASIFICACION_RESOLUCION", "FUENTE_RESOLUCION", "METODO_RESOLUCION", "OBSERVACION"])
    write_csv(AUDITORIAS / "AUDITORIA_VIGENCIA_2025.csv", vigencia, ["CODCLI", "NUM_DOCUMENTO", "CODIGO_UNICO", "VIGENCIA_PROPUESTA", "EVIDENCIA_2025", "ARCHIVO_FUENTE", "ESTADO_VALIDACION", "REQUIERE_CONFIRMACION", "OBSERVACION"])
    write_csv(AUDITORIAS / "NOMINA_SOLICITUD_DATOS_DOCENCIA.csv", simple_manual, ["CODCLI", "NOMBRE_COMPLETO", "CARRERA", "SEDE", "RUN_O_PASAPORTE", "NACIONALIDAD", "TIPO_RESIDENCIA_PENDIENTE", "PAIS_ORIGEN_PENDIENTE", "PAIS_ESTUDIOS_SECUNDARIOS_PENDIENTE", "OTRO_CAMPO_PENDIENTE", "OBSERVACION"])
    write_csv(AUDITORIAS / "NOMINA_GESTION_MANUAL_DETALLE_FASE2.csv", detail_manual, ["ID_CONTROL", "CODCLI", "NOMBRE_COMPLETO", "CARRERA", "SEDE", "CAMPO", "ESTADO_CAMPO", "MOTIVO", "FUENTE_REVISADA"])
    write_csv(AUDITORIAS / "COBERTURA_VARIABLES_BASE_MAESTRA.csv", coverage, ["VARIABLE", "TOTAL_REGISTROS", "COMPLETOS", "PENDIENTES", "CONFLICTOS", "NO_APLICA", "PORCENTAJE_COBERTURA", "FUENTE_PRINCIPAL", "ACCION_PENDIENTE"])
    write_csv(AUDITORIAS / "CATALOGO_PAISES_SIES_2026.csv", countries, ["CODIGO_PAIS", "NOMBRE_PAIS", "ETIQUETA", "PERMITIDO_NACIONALIDAD", "PERMITIDO_PAIS_ORIGEN", "FUENTE"])
    create_management_workbook(GESTION / "PLANILLA_GESTION_DOCENCIA_EXTRANJEROS_2025.xlsx", base_rows, countries, coverage, resolution)
    shutil.copy2(GESTION / "PLANILLA_GESTION_DOCENCIA_EXTRANJEROS_2025.xlsx", GESTION / "PRUEBA_PLANILLA_GESTION_DOCENCIA_EXTRANJEROS_2025_SIN_RESPUESTAS.xlsx")
    outputs.append(GESTION / "PRUEBA_PLANILLA_GESTION_DOCENCIA_EXTRANJEROS_2025_SIN_RESPUESTAS.xlsx")

    write_text(REPORTES / "REPORTE_BASE_MAESTRA_FASE2.md", report_text(base_rows, universe, coverage, resolution, vigencia, backup_dir))
    write_csv(AUDITORIAS / "MANIFIESTO_ARCHIVOS_FASE2.csv", manifest(outputs, backup_dir), ["RUTA_RELATIVA", "NOMBRE", "TAMANO_BYTES", "SHA256", "FECHA_GENERACION", "ORIGEN", "OBSERVACION"])

    print(json.dumps({"universo": Counter(r["CLASIFICACION_UNIVERSO"] for r in universe), "base_rows": len(base_rows), "gestion_manual": len(simple_manual), "backup": rel(backup_dir)}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
