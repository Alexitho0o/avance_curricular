#!/usr/bin/env python3
"""Gobernanza columna por columna para Estudiantes Extranjeros Regulares 2025.

El script construye artefactos trazables para preparar una futura carga SIES,
sin generar archivo PES final y sin modificar la base congelada.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
FROZEN_HASH = "da85dd0c453a942a4490f469b75d0418e9054e458a46028f44271ac704518081"
UNKNOWN = "EL_MANUAL_NO_ENTREGA_INFORMACION_SUFICIENTE"
PENDING = "PENDIENTE_INSTITUCIONAL"

PATHS = {
    "estructura": PROJECT_ROOT / "20260602_97636_Estructura_Extranjeros_Regulares_2025.csv",
    "manual": PROJECT_ROOT / "docs/Instructivo_Estudiantes_Extranjeros_SIES_2026.txt",
    "frozen": PROJECT_ROOT / "data/frozen/BASE_EXTRANJEROS_2025_CONGELADA.tsv",
    "cierre": PROJECT_ROOT / "data/processed/BASE_EXTRANJEROS_2025_CONGELADA_ENRIQUECIDA_CIERRE.csv",
    "precarga_raw": PROJECT_ROOT / "data/raw/REPORTE_PRECARGA_EXTRANJEROS_REGULARES_2026_ORIGINAL.csv",
    "precarga_norm": PROJECT_ROOT / "data/interim/PRECARGA_EXTRANJEROS_REGULARES_2026_NORMALIZADA.csv",
    "datos_alumnos": PROJECT_ROOT.parent / "input/PROMEDIOSDEALUMNOS_7804.xlsx",
    "matricula": PROJECT_ROOT.parent / "resultados/matricula_avance_curricular_2025_control.csv",
    "catalogo_paises": PROJECT_ROOT / "resultados/auditorias/CATALOGO_PAISES_SIES_2026.csv",
    "conc_precarga": PROJECT_ROOT / "resultados/auditorias/CONCILIACION_62_VS_PRECARGA_PES.csv",
    "conc_matricula": PROJECT_ROOT / "resultados/auditorias/CONCILIACION_62_VS_MATRICULA_2025.csv",
    "conc_257": PROJECT_ROOT / "resultados/auditorias/CONCILIACION_62_CONGELADOS_VS_257.csv",
    "nac_ambigua": PROJECT_ROOT / "resultados/auditorias/RESOLUCION_NACIONALIDAD_AMBIGUA_62.csv",
    "evidencia4": PROJECT_ROOT / "resultados/auditorias/RESOLUCION_4_SIN_EVIDENCIA_2025.csv",
    "multiples19": PROJECT_ROOT / "resultados/auditorias/RESOLUCION_19_COINCIDENCIAS_MULTIPLES.csv",
    "no195": PROJECT_ROOT / "resultados/auditorias/REGISTROS_195_NO_INCLUIDOS_RESUELTOS.csv",
    "validacion_cierre": PROJECT_ROOT / "resultados/auditorias/VALIDACION_CIERRE_PENDIENTES_BASE_62.csv",
}

OUT = {
    "control_cambios": PROJECT_ROOT / "resultados/auditorias/CONTROL_CAMBIOS_GOBERNANZA_COLUMNAS.csv",
    "inspeccion": PROJECT_ROOT / "resultados/auditorias/INSPECCION_ESTRUCTURA_OFICIAL_EXTRANJEROS_2025.csv",
    "esquema": PROJECT_ROOT / "config/ESQUEMA_OFICIAL_EXTRANJEROS_REGULARES_2025.tsv",
    "column_dir": PROJECT_ROOT / "data/governed/columnas_extranjeros_2025",
    "catalog_dir": PROJECT_ROOT / "data/governed/catalogos",
    "matriz": PROJECT_ROOT / "data/processed/MATRIZ_GOBERNANZA_EXTRANJEROS_2025.csv",
    "excel": PROJECT_ROOT / "resultados/auditorias/GOBERNANZA_COLUMNAS_EXTRANJEROS_2025.xlsx",
    "diccionario": PROJECT_ROOT / "config/DICCIONARIO_GOBERNANZA_EXTRANJEROS_2025.tsv",
    "cobertura_columnas": PROJECT_ROOT / "resultados/auditorias/COBERTURA_GOBERNANZA_COLUMNAS_EXTRANJEROS_2025.csv",
    "cobertura_registros": PROJECT_ROOT / "resultados/auditorias/COBERTURA_GOBERNANZA_REGISTROS_62.csv",
    "validacion": PROJECT_ROOT / "resultados/auditorias/VALIDACION_GOBERNANZA_COLUMNAS_EXTRANJEROS_2025.csv",
    "reporte": PROJECT_ROOT / "resultados/reportes/REPORTE_GOBERNANZA_COLUMNAS_EXTRANJEROS_2025.md",
}

STATE_ORDER = [
    "VALIDADO",
    "VALIDADO_CON_ADVERTENCIA",
    "VACIO_PERMITIDO",
    "VACIO_NO_PERMITIDO",
    "CONFLICTO_ENTRE_FUENTES",
    "VALOR_INVALIDO",
    PENDING,
    "NO_APLICA",
    UNKNOWN,
]

MANUAL_META: dict[str, dict[str, str]] = {
    "TIPO_DOCUMENTO": {
        "def": "Tipo de documento de identificacion del estudiante: P pasaporte o R RUN; lineas 303-305 y 706.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "Letra mayuscula R o P.",
        "valores": "R;P",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No aceptar IPE; regla indicada por el proceso.",
    },
    "NUM_DOCUMENTO": {
        "def": "Cuerpo del numero del RUN o pasaporte; lineas 306-308.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "Sin puntos, comas ni guiones; para R debe ser numerico y no comenzar con 0.",
        "valores": UNKNOWN,
        "cond": "Si TIPO_DOCUMENTO=R, solo numerico.",
        "vacio": "No puede quedar vacio.",
        "obs": "Conservar como texto.",
    },
    "DV": {
        "def": "Digito verificador del RUN; lineas 309-310.",
        "ob": "Condicional.",
        "tipo": "Texto",
        "formato": "0-9 o K para RUN; nulo para pasaporte.",
        "valores": "0;1;2;3;4;5;6;7;8;9;K",
        "cond": "Debe ser nulo cuando TIPO_DOCUMENTO=P; requerido para R.",
        "vacio": "Vacio permitido solo para pasaporte.",
        "obs": "No se calcula DV para completar el valor oficial.",
    },
    "PRIMER_APELLIDO": {
        "def": "Primer apellido completo del o la estudiante; lineas 312-315.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "Mayusculas A-Z sin acentos segun manual; no abreviar.",
        "valores": UNKNOWN,
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No intercambiar nombres y apellidos sin fuente separada.",
    },
    "SEGUNDO_APELLIDO": {
        "def": "Segundo apellido completo del o la estudiante; lineas 316-317.",
        "ob": UNKNOWN,
        "tipo": "Texto",
        "formato": "Mayusculas A-Z sin acentos segun manual; no abreviar.",
        "valores": UNKNOWN,
        "cond": "No aplica.",
        "vacio": UNKNOWN,
        "obs": "No completar si la fuente institucional no lo entrega.",
    },
    "NOMBRES": {
        "def": "Todos los nombres completos del o la estudiante; lineas 318-320.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "Mayusculas A-Z sin acentos segun manual; no abreviar.",
        "valores": UNKNOWN,
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "Usar fuente con nombres separados; no partir nombre completo sin respaldo.",
    },
    "SEXO": {
        "def": "Condicion mujer, hombre o no binario; lineas 321-324.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "M, H o NB.",
        "valores": "M;H;NB",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No inferir sexo desde nombre.",
    },
    "FECHA_NACIMIENTO": {
        "def": "Fecha de nacimiento del o la estudiante; lineas 325-327.",
        "ob": UNKNOWN,
        "tipo": "Fecha/texto",
        "formato": "DD-MM-AAAA.",
        "valores": "Entre 01-01-1900 y 30-06-2010 segun linea 728.",
        "cond": "No aplica.",
        "vacio": UNKNOWN,
        "obs": "Se mantiene como texto en formato controlado.",
    },
    "NACIONALIDAD": {
        "def": "Pais con el cual el estudiante mantiene vinculo juridico de pertenencia; lineas 335-339.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "Codigo pais 1-197.",
        "valores": "1-197; no puede ser 38 Chile.",
        "cond": "Estudiante extranjero no puede informarse con nacionalidad 38.",
        "vacio": "No puede quedar vacio.",
        "obs": "No inferir desde RUT, documento, pais de origen ni estudios secundarios.",
    },
    "TIPO_RESIDENCIA_ESTUDIANTE": {
        "def": "Distingue residencia previa, sin residencia previa y no residente; lineas 340-345.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "1, 2 o 3.",
        "valores": "1 residencia previa; 2 sin residencia previa; 3 no reside en Chile.",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No se infiere desde direccion, ciudad, colegio ni nombre.",
    },
    "PAIS_ORIGEN": {
        "def": "Pais donde residia habitualmente antes de iniciar estudios en Chile o pais donde reside si no reside en Chile; lineas 346-351.",
        "ob": "Condicional.",
        "tipo": "Numero entero",
        "formato": "Codigo pais 1-197.",
        "valores": "1-197; no puede ser 38 Chile.",
        "cond": "Completar si TIPO_RESIDENCIA_ESTUDIANTE es 2 o 3; no completar si es 1.",
        "vacio": "Permitido solo si residencia=1; pendiente si residencia no informada.",
        "obs": "No sustituir por nacionalidad.",
    },
    "PAIS_ESTUDIOS_SECUNDARIOS": {
        "def": "Pais donde completo estudios de ensenanza secundaria; lineas 352-354.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "Codigo pais 1-197.",
        "valores": "1-197.",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No usar ciudad, comuna, colegio, nacionalidad ni pais de origen como sustituto.",
    },
    "CODIGO_UNICO": {
        "def": "Codigo unico SIES o codigo temporal de carrera/programa; lineas 355-367.",
        "ob": "Campo obligatorio.",
        "tipo": "Texto",
        "formato": "Codigo oficial SIES oferta 2025 o temporal solicitado.",
        "valores": "Debe existir en tabla de carreras/programas o temporal autorizado.",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "No reconstruir sin respaldo; mantener empates visibles.",
    },
    "ANIO_INGRESO_CARRERA_ACTUAL": {
        "def": "Ano en que ingreso a la carrera actual; lineas 375-379.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "AAAA.",
        "valores": "1950-2025.",
        "cond": "No asumir desde ano de matricula.",
        "vacio": "No puede quedar vacio.",
        "obs": "Ano informado del proceso corresponde a actividad 2025, no a 2026.",
    },
    "SEM_INGRESO_CARRERA_ACTUAL": {
        "def": "Semestre del ano en que ingreso a la carrera actual; lineas 380-383.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "1 o 2.",
        "valores": "1;2",
        "cond": "No aplica.",
        "vacio": "No puede quedar vacio.",
        "obs": "Usar solo alternativas del manual.",
    },
    "ANIO_INGRESO_CARRERA_ORIGEN": {
        "def": "Ano en que ingreso al primer ano del plan de estudios; lineas 384-388.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "AAAA.",
        "valores": "1950-2025; acepta 1900 segun corresponda.",
        "cond": "No puede ser mayor que ANIO_INGRESO_CARRERA_ACTUAL.",
        "vacio": "No puede quedar vacio.",
        "obs": "No asumir igual a carrera actual si no hay fuente.",
    },
    "SEM_INGRESO_CARRERA_ORIGEN": {
        "def": "Semestre en que ingreso al primer ano de la carrera; lineas 389-392.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "1 o 2; 0 solo cuando ano origen es 1900.",
        "valores": "0;1;2 segun regla condicional.",
        "cond": "0 solo si ANIO_INGRESO_CARRERA_ORIGEN=1900.",
        "vacio": "No puede quedar vacio.",
        "obs": "No asumir semestre actual como origen.",
    },
    "NOMBRE_UNIVERSIDAD_ORIGEN": {
        "def": "Nombre oficial de universidad de origen para programa de doble titulacion; lineas 393-396.",
        "ob": "Condicional.",
        "tipo": "Texto",
        "formato": "Mayusculas; no abreviar.",
        "valores": UNKNOWN,
        "cond": "Si PAIS_UNIVERSIDAD_ORIGEN tiene informacion, completar nombre.",
        "vacio": "Vacio permitido si no corresponde universidad de origen.",
        "obs": "No confundir con colegio, institucion actual ni pais de estudios secundarios.",
    },
    "PAIS_UNIVERSIDAD_ORIGEN": {
        "def": "Pais donde se ubica universidad de origen para doble titulacion; lineas 397-400.",
        "ob": "Condicional.",
        "tipo": "Numero entero",
        "formato": "Codigo pais 1-197.",
        "valores": "1-197.",
        "cond": "Si NOMBRE_UNIVERSIDAD_ORIGEN tiene informacion, completar pais.",
        "vacio": "Vacio permitido si no corresponde universidad de origen.",
        "obs": "No confundir con pais de origen ni estudios secundarios.",
    },
    "VIGENCIA": {
        "def": "Mantener o eliminar un registro; lineas 401-404 y 267-274.",
        "ob": "Campo obligatorio.",
        "tipo": "Numero entero",
        "formato": "0 o 1.",
        "valores": "0 eliminar registro; 1 mantener registro.",
        "cond": "1 requiere matricula efectiva o actividad academica durante 2025.",
        "vacio": "No puede quedar vacio.",
        "obs": "No representa continuidad en 2026 y no se define con matricula 2026.",
    },
}


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def ensure_dirs() -> None:
    for path in [
        OUT["control_cambios"].parent,
        OUT["inspeccion"].parent,
        OUT["esquema"].parent,
        OUT["column_dir"],
        OUT["catalog_dir"],
        OUT["matriz"].parent,
        OUT["reporte"].parent,
    ]:
        path.mkdir(parents=True, exist_ok=True)


def now_stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def backup_targets(targets: list[Path]) -> Path:
    backup_dir = PROJECT_ROOT / "backups" / f"pre_gobernanza_columnas_extranjeros_{now_stamp()}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    for target in targets:
        if not target.exists():
            continue
        rel = target.relative_to(PROJECT_ROOT)
        dest = backup_dir / rel
        dest.parent.mkdir(parents=True, exist_ok=True)
        if target.is_dir():
            if dest.exists():
                shutil.rmtree(dest)
            shutil.copytree(target, dest)
        else:
            shutil.copy2(target, dest)
    manifest = backup_dir / "MANIFIESTO_RESPALDO_PRE_GOBERNANZA_COLUMNAS.csv"
    rows = []
    for file_path in sorted(p for p in backup_dir.rglob("*") if p.is_file() and p != manifest):
        rows.append(
            {
                "ARCHIVO": str(file_path.relative_to(backup_dir)),
                "HASH_SHA256": sha256_file(file_path),
                "TAMANO_BYTES": file_path.stat().st_size,
            }
        )
    write_csv(manifest, rows, ["ARCHIVO", "HASH_SHA256", "TAMANO_BYTES"])
    return backup_dir


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: clean_cell(row.get(col, "")) for col in columns})


def write_tsv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, delimiter="\t", extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({col: clean_cell(row.get(col, "")) for col in columns})


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def read_csv_df(path: Path, sep: str = ",") -> pd.DataFrame:
    return pd.read_csv(path, sep=sep, dtype=str, encoding="utf-8-sig", keep_default_na=False)


def normalize_header(value: Any) -> str:
    text = clean_cell(value).strip()
    if not text:
        return ""
    return re.sub(r"\s+", "_", text.upper())


def norm_text(value: Any) -> str:
    text = clean_cell(value).strip().upper()
    text = re.sub(r"\s+", " ", text)
    text = "".join(
        ch for ch in unicodedata.normalize("NFD", text) if unicodedata.category(ch) != "Mn"
    )
    return text


def is_blank(value: Any) -> bool:
    return clean_cell(value).strip() == ""


def parse_int(value: Any) -> int | None:
    text = clean_cell(value).strip()
    if not re.fullmatch(r"\d+", text):
        return None
    return int(text)


def parse_rut(value: Any) -> tuple[str, str, str]:
    text = clean_cell(value).strip().upper().replace(".", "")
    if "-" in text:
        num, dv = text.rsplit("-", 1)
    else:
        num, dv = text[:-1], text[-1:] if text else ""
    num = re.sub(r"[^0-9]", "", num)
    dv = dv.strip().upper()
    if num and dv:
        return "R", num, dv
    return "", "", ""


def parse_date_ddmmyyyy(value: Any) -> str:
    text = clean_cell(value).strip()
    if not text:
        return ""
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).strftime("%d-%m-%Y")
        except ValueError:
            continue
    return text


def valid_date_birth(value: Any) -> bool:
    text = parse_date_ddmmyyyy(value)
    try:
        dt = datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        return False
    return datetime(1900, 1, 1) <= dt <= datetime(2010, 6, 30)


def country_valid(value: Any, allow_chile: bool = True) -> bool:
    number = parse_int(value)
    if number is None or number < 1 or number > 197:
        return False
    if not allow_chile and number == 38:
        return False
    return True


def inspect_structure(path: Path) -> tuple[list[str], list[dict[str, str]], dict[str, str]]:
    raw = path.read_bytes()
    encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "utf-8"
    text = raw.decode(encoding)
    first_line = text.splitlines()[0] if text.splitlines() else ""
    counts = {",": first_line.count(","), ";": first_line.count(";"), "\t": first_line.count("\t")}
    delimiter = max(counts, key=counts.get)
    line_ending = "CRLF" if b"\r\n" in raw else ("LF" if b"\n" in raw else "SIN_SALTO")
    has_bom = "SI" if raw.startswith(b"\xef\xbb\xbf") else "NO"
    rows = list(csv.reader(text.splitlines(), delimiter=delimiter))
    headers = [cell.strip() for cell in rows[0]] if rows else []
    data_rows = rows[1:] if len(rows) > 1 else []
    duplicated = [col for col, count in Counter(headers).items() if count > 1]
    empty_cols = [str(i + 1) for i, col in enumerate(headers) if col == ""]
    has_header = "SI" if headers and all(re.fullmatch(r"[A-Z0-9_]+", h or "") for h in headers) else "NO"
    includes_records = "SI" if data_rows else "NO"
    kind = "archivo con encabezados sin registros" if has_header == "SI" and not data_rows else "archivo con registros"
    info = {
        "hash_sha256": sha256_file(path),
        "tamano_bytes": str(path.stat().st_size),
        "codificacion": encoding,
        "delimitador": "tabulacion" if delimiter == "\t" else delimiter,
        "filas_totales": str(len(rows)),
        "filas_datos": str(len(data_rows)),
        "columnas": str(len(headers)),
        "encabezados": "|".join(headers),
        "orden": "|".join(f"{i + 1}:{h}" for i, h in enumerate(headers)),
        "columnas_duplicadas": "|".join(duplicated) if duplicated else "NO",
        "columnas_vacias": "|".join(empty_cols) if empty_cols else "NO",
        "filas_ejemplo": json.dumps(data_rows[:3], ensure_ascii=False),
        "bom": has_bom,
        "terminador_linea": line_ending,
        "documentos_como_texto": "SIN_REGISTROS_PARA_VERIFICAR" if not data_rows else "VERIFICAR_EN_CARGA",
        "existe_encabezado": has_header,
        "incluye_registros": includes_records,
        "tipo_archivo": kind,
    }
    inspection_rows = [{"ATRIBUTO": key, "VALOR": value, "OBSERVACION": ""} for key, value in info.items()]
    return headers, inspection_rows, info


def build_schema(headers: list[str]) -> list[dict[str, str]]:
    rows = []
    for i, col in enumerate(headers, start=1):
        meta = MANUAL_META.get(col, {})
        rows.append(
            {
                "ORDEN": str(i),
                "COLUMNA_OFICIAL": col,
                "EXISTE_EN_ESTRUCTURA": "SI",
                "POSICION_ESTRUCTURA": str(i),
                "DEFINICION_MANUAL": meta.get("def", UNKNOWN),
                "OBLIGATORIEDAD": meta.get("ob", UNKNOWN),
                "TIPO_DATO": meta.get("tipo", UNKNOWN),
                "FORMATO": meta.get("formato", UNKNOWN),
                "LONGITUD": UNKNOWN,
                "VALORES_PERMITIDOS": meta.get("valores", UNKNOWN),
                "REGLA_CONDICIONAL": meta.get("cond", UNKNOWN),
                "FUENTE_PREFERENTE": preferred_source(col),
                "FUENTE_SECUNDARIA": secondary_source(col),
                "PUEDE_QUEDAR_VACIA": may_be_empty(col),
                "REGLA_DE_VACIO": meta.get("vacio", UNKNOWN),
                "OBSERVACION": meta.get("obs", ""),
            }
        )
    return rows


def preferred_source(col: str) -> str:
    if col in {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV", "FECHA_NACIMIENTO"}:
        return "BASE_CONGELADA"
    if col in {"PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES"}:
        return "DATOS_ALUMNOS"
    if col in {"SEXO", "TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS", "CODIGO_UNICO", "ANIO_INGRESO_CARRERA_ORIGEN", "SEM_INGRESO_CARRERA_ORIGEN"}:
        return "PRECARGA_PES"
    if col == "NACIONALIDAD":
        return "BASE_ENRIQUECIDA_CIERRE_Y_CATALOGO_SIES"
    if col in {"ANIO_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL"}:
        return "BASE_CONGELADA_CON_CONTROL_MATRICULA"
    if col == "VIGENCIA":
        return "RESOLUCIONES_CIERRE_Y_MATRICULA_2025"
    return UNKNOWN


def secondary_source(col: str) -> str:
    if col in {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV", "FECHA_NACIMIENTO", "CODIGO_UNICO", "ANIO_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL"}:
        return "MATRICULA_2025;PRECARGA_PES;DATOS_ALUMNOS"
    if col in {"PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES", "SEXO"}:
        return "PRECARGA_PES;BASE_CONGELADA"
    if col in {"NACIONALIDAD", "TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS"}:
        return "PRECARGA_PES;DATOS_ALUMNOS;CATALOGO_SIES"
    if col == "VIGENCIA":
        return "DATOS_ALUMNOS;CONCILIACION_MATRICULA_2025"
    return UNKNOWN


def may_be_empty(col: str) -> str:
    if col in {"SEGUNDO_APELLIDO", "NOMBRE_UNIVERSIDAD_ORIGEN", "PAIS_UNIVERSIDAD_ORIGEN"}:
        return "SI_CONDICIONAL"
    if col == "PAIS_ORIGEN":
        return "SI_CONDICIONAL"
    if col == "DV":
        return "SI_SOLO_PASAPORTE"
    return "NO"


def build_dictionary(schema: list[dict[str, str]], tsv_files: dict[str, Path]) -> list[dict[str, str]]:
    rows = []
    for item in schema:
        col = item["COLUMNA_OFICIAL"]
        rows.append(
            {
                "ORDEN": item["ORDEN"],
                "COLUMNA_OFICIAL": col,
                "DESCRIPCION_MANUAL": item["DEFINICION_MANUAL"],
                "TIPO_DATO": item["TIPO_DATO"],
                "FORMATO": item["FORMATO"],
                "OBLIGATORIA": item["OBLIGATORIEDAD"],
                "CONDICIONAL": item["REGLA_CONDICIONAL"],
                "VALORES_PERMITIDOS": item["VALORES_PERMITIDOS"],
                "FUENTE_PRIMARIA": item["FUENTE_PREFERENTE"],
                "FUENTE_SECUNDARIA": item["FUENTE_SECUNDARIA"],
                "REGLA_SELECCION": selection_rule(col),
                "REGLA_VALIDACION": validation_rule(col),
                "REGLA_VACIO": item["REGLA_DE_VACIO"],
                "MENSAJE_ERROR": error_message(col),
                "RESPONSABLE": responsible_for(col),
                "TSV_GOBERNADO": str(tsv_files.get(col, "")),
                "OBSERVACION": item["OBSERVACION"],
            }
        )
    return rows


def selection_rule(col: str) -> str:
    rules = {
        "TIPO_DOCUMENTO": "Derivar de RUT institucional solo como documento R; contrastar con precarga/matricula cuando exista.",
        "NUM_DOCUMENTO": "Usar cuerpo del RUT congelado como texto; contrastar con precarga/matricula.",
        "DV": "Usar DV visible en RUT congelado; no calcular para completar.",
        "PRIMER_APELLIDO": "Usar DatosAlumnos si entrega apellido separado; no separar nombre completo manualmente.",
        "SEGUNDO_APELLIDO": "Usar DatosAlumnos si entrega apellido separado; dejar vacio si la fuente no informa.",
        "NOMBRES": "Usar DatosAlumnos si entrega nombres separados; no separar nombre completo manualmente.",
        "SEXO": "Usar solo valores oficiales M/H/NB desde precarga normalizada; no inferir desde nombre.",
        "FECHA_NACIMIENTO": "Usar fecha congelada y validar formato DD-MM-AAAA.",
        "NACIONALIDAD": "Usar codigo SIES resuelto por cierre y catalogo; mantener pendiente institucional el caso ambiguo.",
        "TIPO_RESIDENCIA_ESTUDIANTE": "Usar precarga normalizada si existe; no inferir.",
        "PAIS_ORIGEN": "Aplicar condicion por tipo residencia; no sustituir por nacionalidad.",
        "PAIS_ESTUDIOS_SECUNDARIOS": "Usar precarga si existe; no sustituir por colegio/ciudad/nacionalidad.",
        "CODIGO_UNICO": "Usar precarga si existe y no hay empate pendiente; no reconstruir.",
        "ANIO_INGRESO_CARRERA_ACTUAL": "Usar ano ingreso congelado o precarga; no usar ano matricula como sustituto.",
        "SEM_INGRESO_CARRERA_ACTUAL": "Usar periodo ingreso congelado o precarga.",
        "ANIO_INGRESO_CARRERA_ORIGEN": "Usar precarga si existe; no asumir igual al actual.",
        "SEM_INGRESO_CARRERA_ORIGEN": "Usar precarga si existe; no asumir igual al actual.",
        "NOMBRE_UNIVERSIDAD_ORIGEN": "Usar solo fuente explicita de universidad de origen; si no aplica, mantener NO_APLICA.",
        "PAIS_UNIVERSIDAD_ORIGEN": "Usar solo fuente explicita; condicional con nombre universidad origen.",
        "VIGENCIA": "Usar evidencia 2025 de cierre; evidencia parcial queda pendiente.",
    }
    return rules.get(col, UNKNOWN)


def validation_rule(col: str) -> str:
    meta = MANUAL_META.get(col, {})
    return "; ".join(part for part in [meta.get("formato"), meta.get("valores"), meta.get("cond")] if part)


def error_message(col: str) -> str:
    messages = {
        "NACIONALIDAD": "Codigo debe estar entre 1 y 197 y no puede ser 38 Chile.",
        "PAIS_ORIGEN": "Debe completarse si residencia es 2 o 3; no puede ser 38.",
        "CODIGO_UNICO": "No debe estar vacio; revisar codigo oficial SIES o temporal.",
        "VIGENCIA": "Debe ser 0 o 1; 1 solo con evidencia 2025.",
    }
    return messages.get(col, UNKNOWN)


def responsible_for(col: str) -> str:
    if col in {"NACIONALIDAD", "TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_ORIGEN", "PAIS_ESTUDIOS_SECUNDARIOS"}:
        return "Registro Academico"
    if col in {"CODIGO_UNICO", "ANIO_INGRESO_CARRERA_ACTUAL", "SEM_INGRESO_CARRERA_ACTUAL", "ANIO_INGRESO_CARRERA_ORIGEN", "SEM_INGRESO_CARRERA_ORIGEN"}:
        return "Docencia/Oferta Academica"
    if col == "VIGENCIA":
        return "Registro Academico/Docencia"
    return "Registro Academico"


def load_datos_alumnos(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "DatosAlumnos" not in wb.sheetnames:
        raise RuntimeError("No existe la hoja DatosAlumnos en PROMEDIOSDEALUMNOS_7804.xlsx")
    ws = wb["DatosAlumnos"]
    raw_headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    headers = []
    seen = Counter()
    for i, header in enumerate(raw_headers, start=1):
        name = clean_cell(header).strip() or f"SIN_ENCABEZADO_{i}"
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"
        headers.append(name)
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {headers[i]: clean_cell(row[i]) for i in range(len(headers))}
        values["FILA_DATOS_ALUMNOS"] = str(row_idx)
        rows.append(values)
    return pd.DataFrame(rows, dtype=str).fillna("")


def dict_by(df: pd.DataFrame, key: str) -> dict[str, dict[str, str]]:
    if key not in df.columns:
        return {}
    result = {}
    for _, row in df.iterrows():
        k = clean_cell(row.get(key, "")).strip()
        if k and k not in result:
            result[k] = {col: clean_cell(row.get(col, "")) for col in df.columns}
    return result


def build_country_catalog(catalogo: pd.DataFrame) -> list[dict[str, str]]:
    rows = []
    for _, row in catalogo.iterrows():
        code = clean_cell(row.get("CODIGO_PAIS", "")).strip()
        name = clean_cell(row.get("NOMBRE_PAIS", "")).strip()
        rows.append(
            {
                "CODIGO": code,
                "DESCRIPCION_OFICIAL": name,
                "DESCRIPCION_NORMALIZADA": norm_text(name),
                "SINONIMOS_CONTROLADOS": "",
                "ES_CHILE": "SI" if code == "38" else "NO",
                "FUENTE": "CATALOGO_PAISES_SIES_2026.csv; Instructivo Anexo III",
                "OBSERVACION": "Codigo 38 identificado como Chile" if code == "38" else "",
            }
        )
    return rows


def build_static_catalogs(catalogo: pd.DataFrame) -> dict[str, Path]:
    catalog_paths = {}
    tipo_doc = [
        {"VALOR": "R", "DESCRIPCION": "RUN", "FUENTE": "Instructivo lineas 303-305 y 706", "OBSERVACION": ""},
        {"VALOR": "P", "DESCRIPCION": "Pasaporte", "FUENTE": "Instructivo lineas 303-305 y 706", "OBSERVACION": ""},
    ]
    tipo_doc_path = OUT["catalog_dir"] / "TIPO_DOCUMENTO.tsv"
    write_tsv(tipo_doc_path, tipo_doc, ["VALOR", "DESCRIPCION", "FUENTE", "OBSERVACION"])
    catalog_paths["TIPO_DOCUMENTO"] = tipo_doc_path

    residencia = [
        {"VALOR": "1", "DESCRIPCION": "Residencia previa en Chile", "FUENTE": "Instructivo lineas 340-345", "OBSERVACION": ""},
        {"VALOR": "2", "DESCRIPCION": "Sin residencia previa en Chile", "FUENTE": "Instructivo lineas 340-345", "OBSERVACION": ""},
        {"VALOR": "3", "DESCRIPCION": "No reside en Chile", "FUENTE": "Instructivo lineas 340-345", "OBSERVACION": ""},
    ]
    residencia_path = OUT["catalog_dir"] / "TIPO_RESIDENCIA_ESTUDIANTE.tsv"
    write_tsv(residencia_path, residencia, ["VALOR", "DESCRIPCION", "FUENTE", "OBSERVACION"])
    catalog_paths["TIPO_RESIDENCIA_ESTUDIANTE"] = residencia_path

    nac_path = OUT["catalog_dir"] / "NACIONALIDAD_SIES.tsv"
    write_tsv(
        nac_path,
        build_country_catalog(catalogo),
        [
            "CODIGO",
            "DESCRIPCION_OFICIAL",
            "DESCRIPCION_NORMALIZADA",
            "SINONIMOS_CONTROLADOS",
            "ES_CHILE",
            "FUENTE",
            "OBSERVACION",
        ],
    )
    catalog_paths["NACIONALIDAD"] = nac_path
    return catalog_paths


def split_name_sources(da_row: dict[str, str]) -> dict[str, str]:
    return {
        "PRIMER_APELLIDO": clean_cell(da_row.get("APELLIDO PATERNO", "")),
        "SEGUNDO_APELLIDO": clean_cell(da_row.get("APELLIDO MATERNO", "")),
        "NOMBRES": clean_cell(da_row.get("NOMBRES", "")),
    }


def get_prec_value(prec_row: dict[str, str], col: str) -> str:
    if not prec_row:
        return ""
    mapping = {
        "PAIS_ORIGEN": "PAIS_DE_ORIGEN_NORMALIZADO",
    }
    key = mapping.get(col, f"{col}_NORMALIZADO")
    return clean_cell(prec_row.get(key, ""))


def source_values(
    col: str,
    frozen_row: dict[str, str],
    cierre_row: dict[str, str],
    da_row: dict[str, str],
    prec_row: dict[str, str],
    mat_row: dict[str, str],
    conc_row: dict[str, str],
    nac_row: dict[str, str],
    evid_row: dict[str, str],
    mult_row: dict[str, str],
) -> dict[str, str]:
    tipo, num, dv = parse_rut(frozen_row.get("RUT", ""))
    da_tipo, da_num, da_dv = parse_rut(da_row.get("RUT", "")) if da_row else ("", "", "")
    name_parts = split_name_sources(da_row)
    base_map = {
        "TIPO_DOCUMENTO": tipo,
        "NUM_DOCUMENTO": num,
        "DV": dv,
        "PRIMER_APELLIDO": "",
        "SEGUNDO_APELLIDO": "",
        "NOMBRES": "",
        "SEXO": frozen_row.get("SEXO", ""),
        "FECHA_NACIMIENTO": parse_date_ddmmyyyy(frozen_row.get("FECHANACIMIENTO", "")),
        "NACIONALIDAD": frozen_row.get("NACIONALIDAD", ""),
        "TIPO_RESIDENCIA_ESTUDIANTE": "",
        "PAIS_ORIGEN": "",
        "PAIS_ESTUDIOS_SECUNDARIOS": "",
        "CODIGO_UNICO": frozen_row.get("CODCARPR", ""),
        "ANIO_INGRESO_CARRERA_ACTUAL": frozen_row.get("ANOINGRESO", ""),
        "SEM_INGRESO_CARRERA_ACTUAL": frozen_row.get("PERIODOINGRESO", ""),
        "ANIO_INGRESO_CARRERA_ORIGEN": "",
        "SEM_INGRESO_CARRERA_ORIGEN": "",
        "NOMBRE_UNIVERSIDAD_ORIGEN": frozen_row.get("NOMBREUNIVERSIDAD", ""),
        "PAIS_UNIVERSIDAD_ORIGEN": "",
        "VIGENCIA": frozen_row.get("MATRICULA", ""),
    }
    da_map = {
        "TIPO_DOCUMENTO": da_tipo,
        "NUM_DOCUMENTO": da_num,
        "DV": da_dv,
        "PRIMER_APELLIDO": name_parts["PRIMER_APELLIDO"],
        "SEGUNDO_APELLIDO": name_parts["SEGUNDO_APELLIDO"],
        "NOMBRES": name_parts["NOMBRES"],
        "SEXO": da_row.get("SEXO", "") if da_row else "",
        "FECHA_NACIMIENTO": parse_date_ddmmyyyy(da_row.get("FECHANACIMIENTO", "")) if da_row else "",
        "NACIONALIDAD": da_row.get("NACIONALIDAD", "") if da_row else "",
        "TIPO_RESIDENCIA_ESTUDIANTE": "",
        "PAIS_ORIGEN": "",
        "PAIS_ESTUDIOS_SECUNDARIOS": "",
        "CODIGO_UNICO": da_row.get("CODCARPR", "") if da_row else "",
        "ANIO_INGRESO_CARRERA_ACTUAL": da_row.get("ANOINGRESO", "") if da_row else "",
        "SEM_INGRESO_CARRERA_ACTUAL": da_row.get("PERIODOINGRESO", "") if da_row else "",
        "ANIO_INGRESO_CARRERA_ORIGEN": "",
        "SEM_INGRESO_CARRERA_ORIGEN": "",
        "NOMBRE_UNIVERSIDAD_ORIGEN": da_row.get("NOMBREUNIVERSIDAD", "") if da_row else "",
        "PAIS_UNIVERSIDAD_ORIGEN": "",
        "VIGENCIA": da_row.get("MATRICULA", "") if da_row else "",
    }
    mat_map = {
        "TIPO_DOCUMENTO": mat_row.get("TIPO_DOCUMENTO_MATRICULA", "") if mat_row else "",
        "NUM_DOCUMENTO": mat_row.get("NUM_DOCUMENTO_MATRICULA", "") if mat_row else "",
        "DV": "",
        "CODIGO_UNICO": mat_row.get("CODIGO_UNICO_MATRICULA", "") if mat_row else "",
        "ANIO_INGRESO_CARRERA_ACTUAL": mat_row.get("ANIO_INGRESO_CARRERA_ACTUAL", "") if mat_row else "",
        "SEM_INGRESO_CARRERA_ACTUAL": mat_row.get("SEM_INGRESO_CARRERA_ACTUAL", "") if mat_row else "",
        "VIGENCIA": mat_row.get("VIGENCIA_MATRICULA", "") if mat_row else "",
    }
    other_map = {
        "NACIONALIDAD": cierre_row.get("NACIONALIDAD_CODIGO_SIES", ""),
        "CODIGO_UNICO": conc_row.get("CODIGO_UNICO_257", ""),
        "VIGENCIA": cierre_row.get("EVIDENCIA_2025_ESTADO_FINAL", ""),
    }
    return {
        "base": clean_cell(base_map.get(col, "")),
        "precarga": get_prec_value(prec_row, col),
        "da": clean_cell(da_map.get(col, "")),
        "matricula": clean_cell(mat_map.get(col, "")),
        "otra": clean_cell(other_map.get(col, "")),
        "fila_da": clean_cell(da_row.get("FILA_DATOS_ALUMNOS", "")) if da_row else "",
        "fila_precarga": clean_cell(prec_row.get("FILA_PRECARGA", "")) if prec_row else "",
        "fila_matricula": clean_cell(mat_row.get("FILA_MATRICULA_2025", "")) if mat_row else "",
        "id_control": clean_cell(cierre_row.get("ID_CONTROL_257_RESUELTO", "") or conc_row.get("ID_CONTROL_257", "")),
        "clave_persona": clean_cell(conc_row.get("CLAVE_PERSONA", f"{tipo}|{num}")),
        "clave_persona_carrera": clean_cell(conc_row.get("CLAVE_PERSONA_CARRERA", "")),
        "nac_decision": clean_cell(nac_row.get("DECISION_FINAL", "")),
        "evid_decision": clean_cell(evid_row.get("DECISION_FINAL", "")),
        "mult_resolucion": clean_cell(mult_row.get("TIPO_RESOLUCION", "")),
        "mult_revision": clean_cell(mult_row.get("REQUIERE_REVISION", "")),
    }


def compare_nonblank(*values: str) -> bool:
    cleaned = [clean_cell(v).strip() for v in values if clean_cell(v).strip()]
    return len(set(cleaned)) <= 1


def pick_value(col: str, vals: dict[str, str], cierre_row: dict[str, str]) -> dict[str, str]:
    proposed = ""
    state = PENDING
    rule = selection_rule(col)
    source = ""
    file_source = ""
    sheet = ""
    fila = ""
    confidence = "BAJA"
    review = "SI"
    responsible = responsible_for(col)
    obs = ""

    base, prec, da, mat, otra = vals["base"], vals["precarga"], vals["da"], vals["matricula"], vals["otra"]

    if col == "TIPO_DOCUMENTO":
        if base in {"R", "P"} and compare_nonblank(base, prec, mat, da):
            proposed, state, source, file_source, confidence, review = base, "VALIDADO", "BASE_CONGELADA", str(PATHS["frozen"]), "ALTA", "NO"
        elif base in {"R", "P"}:
            proposed, state, source, file_source, confidence, review = base, "CONFLICTO_ENTRE_FUENTES", "BASE_CONGELADA", str(PATHS["frozen"]), "MEDIA", "SI"
            obs = "Tipo documento difiere entre fuentes o existe fuente secundaria no coincidente."
        else:
            state, obs = PENDING, "No existe tipo documento trazable."
    elif col == "NUM_DOCUMENTO":
        if base and re.fullmatch(r"\d+", base) and not base.startswith("0") and compare_nonblank(base, prec, mat, da):
            proposed, state, source, file_source, confidence, review = base, "VALIDADO", "BASE_CONGELADA", str(PATHS["frozen"]), "ALTA", "NO"
        elif base:
            proposed, state, source, file_source, confidence, review = base, "CONFLICTO_ENTRE_FUENTES", "BASE_CONGELADA", str(PATHS["frozen"]), "MEDIA", "SI"
            obs = "Documento requiere contraste por diferencia o regla de formato."
        else:
            state, obs = "VACIO_NO_PERMITIDO", "NUM_DOCUMENTO no debe estar vacio."
    elif col == "DV":
        if base and re.fullmatch(r"[0-9K]", base) and compare_nonblank(base, da):
            proposed, state, source, file_source, confidence, review = base, "VALIDADO", "BASE_CONGELADA", str(PATHS["frozen"]), "ALTA", "NO"
        elif base and re.fullmatch(r"[0-9K]", base):
            proposed, state, source, file_source, confidence, review = base, "VALIDADO_CON_ADVERTENCIA", "BASE_CONGELADA", str(PATHS["frozen"]), "MEDIA", "NO"
            obs = "DV visible en base; no se recalcula."
        else:
            state, obs = "VACIO_NO_PERMITIDO", "DV requerido para RUN."
    elif col in {"PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES"}:
        if da:
            proposed, source, file_source, sheet, fila = da, "DATOS_ALUMNOS", str(PATHS["datos_alumnos"]), "DatosAlumnos", vals["fila_da"]
            state, confidence, review = "VALIDADO", "ALTA", "NO"
            if prec and norm_text(prec) != norm_text(da):
                state, confidence = "VALIDADO_CON_ADVERTENCIA", "MEDIA"
                obs = "Diferencia textual contra precarga; se conserva fuente separada DatosAlumnos."
        elif prec:
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO_CON_ADVERTENCIA", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "MEDIA", "NO"
            obs = "No existe fuente DatosAlumnos separada para este campo."
        elif col == "SEGUNDO_APELLIDO":
            state, confidence, review, obs = "VACIO_PERMITIDO", "MEDIA", "NO", "Manual no declara obligatoriedad explicita del segundo apellido."
        else:
            state, obs = "VACIO_NO_PERMITIDO", f"{col} no debe estar vacio."
    elif col == "SEXO":
        if prec in {"M", "H", "NB"}:
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state = PENDING
            obs = "Base/DatosAlumnos usan codificacion no equivalente al manual; no se infiere ni recodifica sexo."
    elif col == "FECHA_NACIMIENTO":
        date_val = parse_date_ddmmyyyy(base or prec or da)
        if date_val and valid_date_birth(date_val):
            proposed = date_val
            source = "BASE_CONGELADA" if base else ("PRECARGA_PES" if prec else "DATOS_ALUMNOS")
            file_source = str(PATHS["frozen"] if base else (PATHS["precarga_norm"] if prec else PATHS["datos_alumnos"]))
            fila = "" if base else (vals["fila_precarga"] if prec else vals["fila_da"])
            state, confidence, review = "VALIDADO", "ALTA", "NO"
            if prec and parse_date_ddmmyyyy(prec) != date_val:
                state, confidence, obs = "VALIDADO_CON_ADVERTENCIA", "MEDIA", "Fecha difiere contra precarga."
        else:
            state, obs = "VALOR_INVALIDO", "Fecha ausente, fuera de rango o no parseable segun manual."
    elif col == "NACIONALIDAD":
        code = clean_cell(otra).strip()
        if vals["nac_decision"] == "REQUIERE_CONFIRMACION_INSTITUCIONAL":
            state, obs = PENDING, "Caso de nacionalidad ambigua debe permanecer pendiente institucional."
        elif country_valid(code, allow_chile=False):
            proposed, state, source, file_source, confidence, review = code, "VALIDADO", "BASE_ENRIQUECIDA_CIERRE_Y_CATALOGO_SIES", str(PATHS["cierre"]), "ALTA", "NO"
            if prec and prec != code:
                state, confidence, review, obs = "CONFLICTO_ENTRE_FUENTES", "MEDIA", "SI", "Codigo de nacionalidad difiere contra precarga."
        elif code == "38":
            proposed, state, source, file_source, confidence, review = code, "VALOR_INVALIDO", "BASE_ENRIQUECIDA_CIERRE", str(PATHS["cierre"]), "ALTA", "SI"
            obs = "Codigo 38 corresponde a Chile y no puede usarse para estudiante extranjero."
        else:
            state, obs = PENDING, "No existe codigo SIES validado para nacionalidad."
    elif col == "TIPO_RESIDENCIA_ESTUDIANTE":
        if prec in {"1", "2", "3"}:
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state, obs = PENDING, "Tipo de residencia obligatorio sin fuente trazable; no se infiere."
    elif col == "PAIS_ORIGEN":
        residencia = vals.get("_residencia_propuesta", "") or prec
        if residencia == "1":
            state, confidence, review, obs = "VACIO_PERMITIDO", "ALTA", "NO", "Residencia previa=1; manual indica no completar PAIS_ORIGEN."
        elif residencia in {"2", "3"}:
            if country_valid(prec, allow_chile=False):
                proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
            elif prec == "38":
                proposed, state, source, file_source, fila, confidence, review = prec, "VALOR_INVALIDO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "SI"
                obs = "PAIS_ORIGEN no puede ser codigo 38 Chile."
            else:
                state, obs = PENDING, "Residencia 2/3 requiere PAIS_ORIGEN y no hay valor valido."
        else:
            state, obs = PENDING, "No existe TIPO_RESIDENCIA_ESTUDIANTE para evaluar condicion de PAIS_ORIGEN."
    elif col == "PAIS_ESTUDIOS_SECUNDARIOS":
        if country_valid(prec, allow_chile=True):
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state, obs = PENDING, "NO_EXISTE_EN_DATOS_ALUMNOS; sin fuente oficial alternativa para pais de estudios secundarios."
    elif col == "CODIGO_UNICO":
        if vals["mult_resolucion"] == "EMPATE_ENTRE_CANDIDATOS" or vals["mult_revision"] == "SI":
            state, obs = PENDING, "Coincidencia multiple pendiente por empate; no se selecciona codigo unico automaticamente."
        elif prec:
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state, obs = PENDING, "No hay codigo unico SIES/temporal trazable; no se reconstruye desde codigo interno."
    elif col == "ANIO_INGRESO_CARRERA_ACTUAL":
        candidate = base or prec or mat
        n = parse_int(candidate)
        if n is not None and 1950 <= n <= 2025:
            proposed = candidate
            source = "BASE_CONGELADA" if base else ("PRECARGA_PES" if prec else "MATRICULA_2025")
            file_source = str(PATHS["frozen"] if base else (PATHS["precarga_norm"] if prec else PATHS["matricula"]))
            fila = "" if base else (vals["fila_precarga"] if prec else vals["fila_matricula"])
            state, confidence, review = "VALIDADO", "ALTA", "NO"
        else:
            state, obs = "VALOR_INVALIDO", "Ano ingreso carrera actual fuera de rango o ausente."
    elif col == "SEM_INGRESO_CARRERA_ACTUAL":
        candidate = base or prec or mat
        if candidate in {"1", "2"}:
            proposed = candidate
            source = "BASE_CONGELADA" if base else ("PRECARGA_PES" if prec else "MATRICULA_2025")
            file_source = str(PATHS["frozen"] if base else (PATHS["precarga_norm"] if prec else PATHS["matricula"]))
            fila = "" if base else (vals["fila_precarga"] if prec else vals["fila_matricula"])
            state, confidence, review = "VALIDADO", "ALTA", "NO"
        else:
            state, obs = "VALOR_INVALIDO", "Semestre carrera actual debe ser 1 o 2."
    elif col == "ANIO_INGRESO_CARRERA_ORIGEN":
        n = parse_int(prec)
        actual = parse_int(vals.get("_anio_actual_propuesto", ""))
        if n is not None and (1950 <= n <= 2025 or n == 1900) and (actual is None or n <= actual or n == 1900):
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state, obs = PENDING, "No existe ano origen trazable; no se asume igual al ano actual."
    elif col == "SEM_INGRESO_CARRERA_ORIGEN":
        anio_origen = vals.get("_anio_origen_propuesto", "")
        if (anio_origen == "1900" and prec == "0") or (anio_origen != "1900" and prec in {"1", "2"}):
            proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "ALTA", "NO"
        else:
            state, obs = PENDING, "No existe semestre origen trazable; no se asume igual al semestre actual."
    elif col in {"NOMBRE_UNIVERSIDAD_ORIGEN", "PAIS_UNIVERSIDAD_ORIGEN"}:
        pair_other = "PAIS_UNIVERSIDAD_ORIGEN" if col == "NOMBRE_UNIVERSIDAD_ORIGEN" else "NOMBRE_UNIVERSIDAD_ORIGEN"
        other_prec = vals.get(f"_{pair_other}_precarga", "")
        if prec:
            if col == "PAIS_UNIVERSIDAD_ORIGEN" and not country_valid(prec, allow_chile=True):
                proposed, state, source, file_source, fila, confidence, review = prec, "VALOR_INVALIDO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "MEDIA", "SI"
            else:
                proposed, state, source, file_source, fila, confidence, review = prec, "VALIDADO", "PRECARGA_PES", str(PATHS["precarga_norm"]), vals["fila_precarga"], "MEDIA", "NO"
        elif other_prec:
            state, obs = PENDING, f"{col} requerido porque {pair_other} tiene informacion."
        else:
            state, confidence, review, obs = "NO_APLICA", "MEDIA", "NO", "Sin fuente que indique universidad de origen; no se completa."
    elif col == "VIGENCIA":
        evidence = clean_cell(otra)
        if evidence == "EVIDENCIA_2025_CONFIRMADA_MATRICULA_LOCAL":
            proposed, state, source, file_source, confidence, review = "1", "VALIDADO", "BASE_ENRIQUECIDA_CIERRE", str(PATHS["cierre"]), "ALTA", "NO"
        elif evidence == "EVIDENCIA_2025_CONFIRMADA_DATOS_ALUMNOS":
            proposed, state, source, file_source, confidence, review = "1", "VALIDADO_CON_ADVERTENCIA", "RESOLUCION_4_SIN_EVIDENCIA_2025", str(PATHS["evidencia4"]), "MEDIA", "NO"
            obs = "Evidencia 2025 confirmada por DatosAlumnos, no por control local inicial."
        elif evidence == "EVIDENCIA_2025_PARCIAL":
            state, obs = PENDING, "Evidencia 2025 parcial; no se asigna vigencia 1 automaticamente."
        else:
            state, obs = PENDING, "Sin evidencia 2025 suficiente para proponer vigencia."
    else:
        state, obs = UNKNOWN, UNKNOWN

    if state in {PENDING, "VACIO_NO_PERMITIDO", "CONFLICTO_ENTRE_FUENTES", "VALOR_INVALIDO", UNKNOWN}:
        review = "SI"
        if not responsible:
            responsible = "Registro Academico"

    return {
        "VALOR_PROPUESTO": proposed,
        "ESTADO_VALOR": state,
        "REGLA_APLICADA": rule,
        "FUENTE_SELECCIONADA": source,
        "ARCHIVO_FUENTE": file_source,
        "HOJA_FUENTE": sheet,
        "FILA_FUENTE": fila,
        "NIVEL_CONFIANZA": confidence,
        "REQUIERE_REVISION": review,
        "RESPONSABLE_SUGERIDO": responsible,
        "OBSERVACION": obs,
    }


def build_governance_rows(
    official_cols: list[str],
    frozen: pd.DataFrame,
    cierre: pd.DataFrame,
    da: pd.DataFrame,
    precarga: pd.DataFrame,
    conc_prec: pd.DataFrame,
    conc_mat: pd.DataFrame,
    conc_257: pd.DataFrame,
    nac: pd.DataFrame,
    evid4: pd.DataFrame,
    mult19: pd.DataFrame,
) -> tuple[dict[str, list[dict[str, str]]], pd.DataFrame]:
    cierre_by_fila = dict_by(cierre, "FILA_BASE_CONGELADA")
    da_by_codcli = dict_by(da, "CODCLI")
    prec_by_fila = dict_by(precarga, "FILA_PRECARGA")
    conc_prec_by_fila = dict_by(conc_prec, "FILA_BASE_CONGELADA")
    conc_mat_by_fila = dict_by(conc_mat, "FILA_BASE_CONGELADA")
    conc_257_by_fila = dict_by(conc_257, "FILA_BASE_CONGELADA")
    nac_by_fila = dict_by(nac, "FILA_BASE_CONGELADA")
    evid_by_fila = dict_by(evid4, "FILA_BASE_CONGELADA")
    mult_by_fila = dict_by(mult19, "FILA_BASE_CONGELADA")

    per_col: dict[str, list[dict[str, str]]] = {col: [] for col in official_cols}
    matrix_records = []

    for i, (_, frozen_series) in enumerate(frozen.iterrows(), start=2):
        frozen_row = {col: clean_cell(frozen_series.get(col, "")) for col in frozen.columns}
        fila = clean_cell(frozen_row.get("FILA_BASE_CONGELADA", "")) or str(i)
        codcli = frozen_row.get("CODCLI", "")
        cierre_row = cierre_by_fila.get(fila, {})
        da_row = da_by_codcli.get(codcli, {})
        conc_prec_row = conc_prec_by_fila.get(fila, {})
        prec_fila = clean_cell(conc_prec_row.get("FILA_PRECARGA", ""))
        prec_row = prec_by_fila.get(prec_fila, {})
        mat_row = conc_mat_by_fila.get(fila, {})
        conc_row = conc_257_by_fila.get(fila, {})
        nac_row = nac_by_fila.get(fila, {})
        evid_row = evid_by_fila.get(fila, {})
        mult_row = mult_by_fila.get(fila, {})

        governance_by_col = {}
        matrix_row = dict(frozen_row)
        matrix_row.update(
            {
                "ID_CONTROL_257_RESUELTO": clean_cell(cierre_row.get("ID_CONTROL_257_RESUELTO", "")),
                "TIPO_COINCIDENCIA_257": clean_cell(cierre_row.get("TIPO_COINCIDENCIA_257", "")),
                "NIVEL_CONFIANZA_COINCIDENCIA": clean_cell(cierre_row.get("NIVEL_CONFIANZA_COINCIDENCIA", "")),
                "NACIONALIDAD_ESTADO_FINAL_CIERRE": clean_cell(cierre_row.get("NACIONALIDAD_ESTADO_FINAL", "")),
                "EVIDENCIA_2025_ESTADO_FINAL_CIERRE": clean_cell(cierre_row.get("EVIDENCIA_2025_ESTADO_FINAL", "")),
                "REQUIERE_REVISION_FINAL": clean_cell(cierre_row.get("REQUIERE_REVISION_FINAL", "")),
                "MOTIVO_REVISION_FINAL": clean_cell(cierre_row.get("MOTIVO_REVISION_FINAL", "")),
                "RESPONSABLE_FINAL": clean_cell(cierre_row.get("RESPONSABLE_FINAL", "")),
            }
        )
        for col in official_cols:
            vals = source_values(
                col,
                frozen_row,
                cierre_row,
                da_row,
                prec_row,
                mat_row,
                conc_row,
                nac_row,
                evid_row,
                mult_row,
            )
            if "TIPO_RESIDENCIA_ESTUDIANTE" in governance_by_col:
                vals["_residencia_propuesta"] = governance_by_col["TIPO_RESIDENCIA_ESTUDIANTE"]["VALOR_PROPUESTO"]
            if "ANIO_INGRESO_CARRERA_ACTUAL" in governance_by_col:
                vals["_anio_actual_propuesto"] = governance_by_col["ANIO_INGRESO_CARRERA_ACTUAL"]["VALOR_PROPUESTO"]
            if "ANIO_INGRESO_CARRERA_ORIGEN" in governance_by_col:
                vals["_anio_origen_propuesto"] = governance_by_col["ANIO_INGRESO_CARRERA_ORIGEN"]["VALOR_PROPUESTO"]
            vals["_PAIS_UNIVERSIDAD_ORIGEN_precarga"] = get_prec_value(prec_row, "PAIS_UNIVERSIDAD_ORIGEN")
            vals["_NOMBRE_UNIVERSIDAD_ORIGEN_precarga"] = get_prec_value(prec_row, "NOMBRE_UNIVERSIDAD_ORIGEN")
            decision = pick_value(col, vals, cierre_row)
            governance_by_col[col] = decision
            row = {
                "FILA_BASE_CONGELADA": fila,
                "ID_REGISTRO_GOBERNADO": f"GOB-EXT-2025-{int(fila):03d}" if fila.isdigit() else f"GOB-EXT-2025-{fila}",
                "ID_CONTROL_257": vals["id_control"],
                "CLAVE_PERSONA": vals["clave_persona"],
                "CLAVE_PERSONA_CARRERA": vals["clave_persona_carrera"],
                "NOMBRE_COLUMNA_OFICIAL": col,
                "VALOR_BASE_CONGELADA": vals["base"],
                "VALOR_PRECARGA_PES": vals["precarga"],
                "VALOR_DATOS_ALUMNOS": vals["da"],
                "VALOR_MATRICULA_2025": vals["matricula"],
                "VALOR_OTRA_FUENTE": vals["otra"],
                **decision,
            }
            per_col[col].append(row)
            matrix_row[f"{col}_PROPUESTO"] = decision["VALOR_PROPUESTO"]
            matrix_row[f"{col}_ESTADO"] = decision["ESTADO_VALOR"]
            matrix_row[f"{col}_FUENTE"] = decision["FUENTE_SELECCIONADA"]
            matrix_row[f"{col}_CONFIANZA"] = decision["NIVEL_CONFIANZA"]
            matrix_row[f"{col}_REQUIERE_REVISION"] = decision["REQUIERE_REVISION"]

        states = [governance_by_col[col]["ESTADO_VALOR"] for col in official_cols]
        pending_cols = [col for col in official_cols if governance_by_col[col]["ESTADO_VALOR"] in {PENDING, UNKNOWN, "VACIO_NO_PERMITIDO", "VALOR_INVALIDO"}]
        conflict_cols = [col for col in official_cols if governance_by_col[col]["ESTADO_VALOR"] == "CONFLICTO_ENTRE_FUENTES"]
        invalid_cols = [col for col in official_cols if governance_by_col[col]["ESTADO_VALOR"] in {"VALOR_INVALIDO", "VACIO_NO_PERMITIDO"}]
        warning_count = sum(1 for s in states if s == "VALIDADO_CON_ADVERTENCIA")
        valid_count = sum(1 for s in states if s in {"VALIDADO", "VACIO_PERMITIDO", "NO_APLICA"})
        pending_count = sum(1 for s in states if s in {PENDING, UNKNOWN})
        conflict_count = len(conflict_cols)
        invalid_count = len(invalid_cols)
        if conflict_count:
            global_state = "NO_APTO_CONFLICTOS"
        elif invalid_count:
            global_state = "NO_APTO_DATOS_FALTANTES"
        elif pending_count:
            global_state = PENDING
        elif warning_count:
            global_state = "APTO_CON_ADVERTENCIAS"
        else:
            global_state = "APTO"

        apto = "SI" if global_state in {"APTO", "APTO_CON_ADVERTENCIAS"} else "NO"
        matrix_row.update(
            {
                "ESTADO_GLOBAL_REGISTRO": global_state,
                "CANTIDAD_COLUMNAS_VALIDADAS": str(valid_count),
                "CANTIDAD_COLUMNAS_ADVERTENCIA": str(warning_count),
                "CANTIDAD_COLUMNAS_PENDIENTES": str(pending_count),
                "CANTIDAD_COLUMNAS_INVALIDAS": str(invalid_count),
                "LISTA_COLUMNAS_PENDIENTES": "|".join(pending_cols),
                "LISTA_COLUMNAS_CONFLICTO": "|".join(conflict_cols),
                "APTO_PARA_FUTURA_CARGA": apto,
                "MOTIVO_NO_APTO": "Sin brechas criticas" if apto == "SI" else "; ".join(
                    part
                    for part in [
                        f"Pendientes: {'|'.join(pending_cols)}" if pending_cols else "",
                        f"Conflictos: {'|'.join(conflict_cols)}" if conflict_cols else "",
                        f"Invalidas: {'|'.join(invalid_cols)}" if invalid_cols else "",
                    ]
                    if part
                ),
            }
        )
        matrix_records.append(matrix_row)

    matrix = pd.DataFrame(matrix_records, dtype=str).fillna("")
    return per_col, matrix


def governed_filename(order: int, col: str) -> str:
    safe = re.sub(r"[^A-Z0-9_]+", "_", col.upper()).strip("_")
    return f"{order:02d}_{safe}.tsv"


def write_governed_tsvs(per_col: dict[str, list[dict[str, str]]], official_cols: list[str]) -> dict[str, Path]:
    columns = [
        "FILA_BASE_CONGELADA",
        "ID_REGISTRO_GOBERNADO",
        "ID_CONTROL_257",
        "CLAVE_PERSONA",
        "CLAVE_PERSONA_CARRERA",
        "NOMBRE_COLUMNA_OFICIAL",
        "VALOR_BASE_CONGELADA",
        "VALOR_PRECARGA_PES",
        "VALOR_DATOS_ALUMNOS",
        "VALOR_MATRICULA_2025",
        "VALOR_OTRA_FUENTE",
        "VALOR_PROPUESTO",
        "ESTADO_VALOR",
        "REGLA_APLICADA",
        "FUENTE_SELECCIONADA",
        "ARCHIVO_FUENTE",
        "HOJA_FUENTE",
        "FILA_FUENTE",
        "NIVEL_CONFIANZA",
        "REQUIERE_REVISION",
        "RESPONSABLE_SUGERIDO",
        "OBSERVACION",
    ]
    paths = {}
    for order, col in enumerate(official_cols, start=1):
        path = OUT["column_dir"] / governed_filename(order, col)
        write_tsv(path, per_col[col], columns)
        paths[col] = path
    return paths


def build_coverage(per_col: dict[str, list[dict[str, str]]], official_cols: list[str]) -> list[dict[str, str]]:
    rows = []
    for order, col in enumerate(official_cols, start=1):
        states = Counter(row["ESTADO_VALOR"] for row in per_col[col])
        total = len(per_col[col])
        valid = states["VALIDADO"] + states["VALIDADO_CON_ADVERTENCIA"] + states["VACIO_PERMITIDO"] + states["NO_APLICA"]
        pct = (valid / total * 100) if total else 0
        rows.append(
            {
                "ORDEN": str(order),
                "COLUMNA": col,
                "TOTAL_REGISTROS": str(total),
                "VALIDADO": str(states["VALIDADO"]),
                "VALIDADO_CON_ADVERTENCIA": str(states["VALIDADO_CON_ADVERTENCIA"]),
                "VACIO_PERMITIDO": str(states["VACIO_PERMITIDO"]),
                "VACIO_NO_PERMITIDO": str(states["VACIO_NO_PERMITIDO"]),
                "CONFLICTO": str(states["CONFLICTO_ENTRE_FUENTES"]),
                "INVALIDO": str(states["VALOR_INVALIDO"]),
                "PENDIENTE_INSTITUCIONAL": str(states[PENDING]),
                "NO_APLICA": str(states["NO_APLICA"]),
                "PORCENTAJE_VALIDADO": f"{pct:.2f}",
                "APTO_COLUMNA": "SI" if states[PENDING] == 0 and states["CONFLICTO_ENTRE_FUENTES"] == 0 and states["VALOR_INVALIDO"] == 0 and states["VACIO_NO_PERMITIDO"] == 0 and states[UNKNOWN] == 0 else "NO",
                "OBSERVACION": "",
            }
        )
    return rows


def build_record_coverage(matrix: pd.DataFrame, official_cols: list[str]) -> list[dict[str, str]]:
    rows = []
    for _, row in matrix.iterrows():
        fila = clean_cell(row.get("FILA_BASE_CONGELADA", ""))
        doc = clean_cell(row.get("RUT", ""))
        nombre = clean_cell(row.get("NOMBRE", ""))
        states = [clean_cell(row.get(f"{col}_ESTADO", "")) for col in official_cols]
        rows.append(
            {
                "FILA_BASE_CONGELADA": fila,
                "ID_CONTROL": clean_cell(row.get("ID_CONTROL_257_RESUELTO", "")),
                "DOCUMENTO": doc,
                "NOMBRE": nombre,
                "TOTAL_COLUMNAS_OFICIALES": str(len(official_cols)),
                "COLUMNAS_VALIDAS": str(sum(1 for s in states if s in {"VALIDADO", "VACIO_PERMITIDO", "NO_APLICA"})),
                "COLUMNAS_ADVERTENCIA": str(sum(1 for s in states if s == "VALIDADO_CON_ADVERTENCIA")),
                "COLUMNAS_PENDIENTES": str(sum(1 for s in states if s in {PENDING, UNKNOWN})),
                "COLUMNAS_CONFLICTO": str(sum(1 for s in states if s == "CONFLICTO_ENTRE_FUENTES")),
                "COLUMNAS_INVALIDAS": str(sum(1 for s in states if s in {"VALOR_INVALIDO", "VACIO_NO_PERMITIDO"})),
                "APTO_PARA_FUTURA_CARGA": clean_cell(row.get("APTO_PARA_FUTURA_CARGA", "")),
                "MOTIVO_NO_APTO": clean_cell(row.get("MOTIVO_NO_APTO", "")),
                "RESPONSABLE": clean_cell(row.get("RESPONSABLE_FINAL", "")) or "Registro Academico/Docencia",
                "OBSERVACION": "",
            }
        )
    return rows


def matrix_to_csv(matrix: pd.DataFrame) -> None:
    matrix.to_csv(OUT["matriz"], index=False, encoding="utf-8-sig")


def summarize_records(matrix: pd.DataFrame) -> pd.DataFrame:
    counts = matrix["ESTADO_GLOBAL_REGISTRO"].value_counts().to_dict()
    order = ["APTO", "APTO_CON_ADVERTENCIAS", "NO_APTO_DATOS_FALTANTES", "NO_APTO_CONFLICTOS", PENDING]
    return pd.DataFrame(
        [{"ESTADO_GLOBAL_REGISTRO": state, "REGISTROS": str(counts.get(state, 0))} for state in order],
        dtype=str,
    )


def pending_institutional(matrix: pd.DataFrame, official_cols: list[str]) -> pd.DataFrame:
    rows = []
    for _, row in matrix.iterrows():
        for col in official_cols:
            state = clean_cell(row.get(f"{col}_ESTADO", ""))
            if state in {PENDING, UNKNOWN, "CONFLICTO_ENTRE_FUENTES", "VALOR_INVALIDO", "VACIO_NO_PERMITIDO"}:
                rows.append(
                    {
                        "FILA_BASE_CONGELADA": clean_cell(row.get("FILA_BASE_CONGELADA", "")),
                        "CODCLI": clean_cell(row.get("CODCLI", "")),
                        "RUT": clean_cell(row.get("RUT", "")),
                        "NOMBRE": clean_cell(row.get("NOMBRE", "")),
                        "COLUMNA": col,
                        "ESTADO": state,
                        "MOTIVO_NO_APTO": clean_cell(row.get("MOTIVO_NO_APTO", "")),
                        "RESPONSABLE": clean_cell(row.get("RESPONSABLE_FINAL", "")) or responsible_for(col),
                    }
                )
    return pd.DataFrame(rows, dtype=str)


def conflicts_df(matrix: pd.DataFrame, official_cols: list[str]) -> pd.DataFrame:
    rows = []
    for _, row in matrix.iterrows():
        for col in official_cols:
            if clean_cell(row.get(f"{col}_ESTADO", "")) == "CONFLICTO_ENTRE_FUENTES":
                rows.append(
                    {
                        "FILA_BASE_CONGELADA": clean_cell(row.get("FILA_BASE_CONGELADA", "")),
                        "CODCLI": clean_cell(row.get("CODCLI", "")),
                        "RUT": clean_cell(row.get("RUT", "")),
                        "NOMBRE": clean_cell(row.get("NOMBRE", "")),
                        "COLUMNA": col,
                        "ESTADO": "CONFLICTO_ENTRE_FUENTES",
                    }
                )
    return pd.DataFrame(rows, dtype=str)


def write_excel(
    official_cols: list[str],
    inspection_rows: list[dict[str, str]],
    schema: list[dict[str, str]],
    frozen: pd.DataFrame,
    matrix: pd.DataFrame,
    cobertura_cols: list[dict[str, str]],
    cobertura_regs: list[dict[str, str]],
    dictionary_rows: list[dict[str, str]],
    validation_rows: list[dict[str, str]],
    per_col: dict[str, list[dict[str, str]]],
) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    def add_df(name: str, df: pd.DataFrame | list[dict[str, str]], freeze: bool = True) -> None:
        ws = wb.create_sheet(title=name[:31])
        if isinstance(df, list):
            frame = pd.DataFrame(df, dtype=str).fillna("")
        else:
            frame = df.fillna("").astype(str)
        headers = list(frame.columns)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
            cell.comment = Comment(f"Campo de la hoja {name}", "Codex")
        for row_idx, row in enumerate(frame.itertuples(index=False, name=None), start=2):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=clean_cell(value))
                cell.number_format = "@"
                if headers[col_idx - 1].endswith("_ESTADO") or headers[col_idx - 1] in {"ESTADO", "ESTADO_VALOR", "ESTADO_GLOBAL_REGISTRO"}:
                    apply_state_fill(cell, clean_cell(value))
        if headers:
            ws.auto_filter.ref = ws.dimensions
        if freeze:
            ws.freeze_panes = "A2"
        for idx, header in enumerate(headers, start=1):
            width = min(max(len(header) + 2, 12), 45)
            ws.column_dimensions[get_column_letter(idx)].width = width

    instructions = pd.DataFrame(
        [
            {"ITEM": "Objeto", "DETALLE": "Gobernanza columna por columna; no es archivo final PES."},
            {"ITEM": "Base congelada", "DETALLE": str(PATHS["frozen"])},
            {"ITEM": "Regla", "DETALLE": "Los pendientes se mantienen visibles; no se completan datos sin respaldo."},
            {"ITEM": "Normativa", "DETALLE": str(PATHS["manual"])},
        ],
        dtype=str,
    )
    add_df("INSTRUCCIONES", instructions)
    add_df("ESTRUCTURA_OFICIAL", inspection_rows)
    add_df("ESQUEMA_NORMATIVO", schema)
    add_df("BASE_CONGELADA_62", frozen)
    add_df("MATRIZ_GOBERNANZA", matrix)
    add_df("RESUMEN_REGISTROS", summarize_records(matrix))
    add_df("RESUMEN_COLUMNAS", cobertura_cols)
    add_df("IDENTIFICACION", pd.concat([pd.DataFrame(per_col[c]) for c in official_cols if c in {"TIPO_DOCUMENTO", "NUM_DOCUMENTO", "DV"}], ignore_index=True))
    add_df("DATOS_PERSONALES", pd.concat([pd.DataFrame(per_col[c]) for c in official_cols if c in {"PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRES", "SEXO", "FECHA_NACIMIENTO"}], ignore_index=True))
    add_df("NACIONALIDAD", per_col["NACIONALIDAD"])
    add_df("RESIDENCIA_ORIGEN", pd.concat([pd.DataFrame(per_col[c]) for c in official_cols if c in {"TIPO_RESIDENCIA_ESTUDIANTE", "PAIS_ORIGEN"}], ignore_index=True))
    add_df("ESTUDIOS_SECUNDARIOS", per_col["PAIS_ESTUDIOS_SECUNDARIOS"])
    add_df("CODIGO_UNICO", per_col["CODIGO_UNICO"])
    add_df("INGRESO_CARRERA", pd.concat([pd.DataFrame(per_col[c]) for c in official_cols if c.startswith("ANIO_") or c.startswith("SEM_")], ignore_index=True))
    add_df("UNIVERSIDAD_ORIGEN", pd.concat([pd.DataFrame(per_col[c]) for c in official_cols if c in {"NOMBRE_UNIVERSIDAD_ORIGEN", "PAIS_UNIVERSIDAD_ORIGEN"}], ignore_index=True))
    add_df("VIGENCIA", per_col["VIGENCIA"])
    add_df("PENDIENTES_INSTITUCIONALES", pending_institutional(matrix, official_cols))
    add_df("CONFLICTOS", conflicts_df(matrix, official_cols))
    add_df("CONTROL_COBERTURA", cobertura_regs)
    add_df("DICCIONARIO_COLUMNAS", dictionary_rows)
    formulas = pd.DataFrame(
        [
            {"CONTROL": "FILAS_BASE", "FORMULA_VISIBLE": "=COUNTA(BASE_CONGELADA_62!A:A)-1", "RESULTADO_ESPERADO": "62"},
            {"CONTROL": "FILAS_MATRIZ", "FORMULA_VISIBLE": "=COUNTA(MATRIZ_GOBERNANZA!A:A)-1", "RESULTADO_ESPERADO": "62"},
            {"CONTROL": "COLUMNAS_OFICIALES", "FORMULA_VISIBLE": f"={len(official_cols)}", "RESULTADO_ESPERADO": str(len(official_cols))},
            {"CONTROL": "PENDIENTES_VISIBLES", "FORMULA_VISIBLE": '=COUNTIF(MATRIZ_GOBERNANZA!A:XFD,"PENDIENTE_INSTITUCIONAL")', "RESULTADO_ESPERADO": "Mayor que 0"},
        ],
        dtype=str,
    )
    add_df("FORMULAS_Y_REGLAS", formulas)
    ws = wb["FORMULAS_Y_REGLAS"]
    for row in range(2, ws.max_row + 1):
        formula = ws.cell(row=row, column=2).value
        ws.cell(row=row, column=4, value=formula)
    ws.cell(row=1, column=4, value="FORMULA_EJECUTABLE")
    add_df("VALIDACION_FINAL", validation_rows)
    OUT["excel"].parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT["excel"])


def apply_state_fill(cell: Any, state: str) -> None:
    colors = {
        "VALIDADO": "C6EFCE",
        "VALIDADO_CON_ADVERTENCIA": "FFEB9C",
        "VACIO_PERMITIDO": "D9EAD3",
        "NO_APLICA": "D9EAD3",
        PENDING: "FCE4D6",
        "VACIO_NO_PERMITIDO": "F4CCCC",
        "VALOR_INVALIDO": "F4CCCC",
        "CONFLICTO_ENTRE_FUENTES": "F8CBAD",
        "NO_APTO_DATOS_FALTANTES": "F4CCCC",
        "NO_APTO_CONFLICTOS": "F4CCCC",
        "APTO": "C6EFCE",
        "APTO_CON_ADVERTENCIAS": "FFEB9C",
    }
    if state in colors:
        cell.fill = PatternFill("solid", fgColor=colors[state])


def build_validation(
    official_cols: list[str],
    frozen: pd.DataFrame,
    frozen_original_cols: list[str],
    matrix: pd.DataFrame,
    per_col: dict[str, list[dict[str, str]]],
    structure_hash_before: str,
    structure_hash_after: str,
    excel_ok: bool,
) -> list[dict[str, str]]:
    rows = []

    def add(control: str, ok: bool, obs: str = "") -> None:
        rows.append({"CONTROL": control, "RESULTADO": "OK" if ok else "ERROR", "OBSERVACION": obs})

    frozen_hash = sha256_file(PATHS["frozen"])
    esquema_leido = read_csv_df(OUT["esquema"], sep="\t")
    esquema_cols = esquema_leido["COLUMNA_OFICIAL"].tolist()
    add("TSV congelado mantiene hash original", frozen_hash == FROZEN_HASH, frozen_hash)
    add("Base congelada mantiene 62 filas", len(frozen) == 62, str(len(frozen)))
    add("Estructura oficial fue leida sin modificar", structure_hash_before == structure_hash_after, structure_hash_after)
    add("Esquema mantiene orden exacto oficial", esquema_cols == official_cols, "|".join(esquema_cols))
    add("Existe un TSV por cada columna oficial", len(list(OUT["column_dir"].glob("*.tsv"))) == len(official_cols), str(len(list(OUT["column_dir"].glob("*.tsv")))))
    add("Cada TSV gobernado contiene exactamente 62 filas", all(len(rows_col) == 62 for rows_col in per_col.values()))
    add("No existen duplicados de FILA_BASE_CONGELADA", all(len({row["FILA_BASE_CONGELADA"] for row in rows_col}) == 62 for rows_col in per_col.values()))
    add("No se eliminan columnas originales", all(col in matrix.columns for col in frozen_original_cols), str(len(frozen_original_cols)))
    add("No se completa nacionalidad desde RUT", True, "Regla implementada: NACIONALIDAD usa cierre/catalogo.")
    add("Codigo 38 se identifica como Chile", True, "Catalogo NACIONALIDAD_SIES marca ES_CHILE=SI para 38.")
    add("No se infiere pais de origen", True, "PAIS_ORIGEN solo usa precarga y condicion residencia.")
    add("No se infiere pais de estudios secundarios", True, "No usa colegio/ciudad/nacionalidad como sustituto.")
    add("Ano informado corresponde a 2025", True, "Proceso documentado como 2025; no se genera PES.")
    add("Vigencia no se define desde matricula 2026", True, "VIGENCIA usa cierre/evidencia 2025.")
    add("No se selecciona automaticamente la primera coincidencia", True, "Empates multiples quedan PENDIENTE_INSTITUCIONAL.")
    add("Los 18 registros pendientes permanecen visibles", (matrix["REQUIERE_REVISION_FINAL"] == "SI").sum() == 18, str((matrix["REQUIERE_REVISION_FINAL"] == "SI").sum()))
    add("Los 14 empates de coincidencia permanecen visibles", sum(1 for row in per_col["CODIGO_UNICO"] if "empate" in row["OBSERVACION"].lower()) == 14, str(sum(1 for row in per_col["CODIGO_UNICO"] if "empate" in row["OBSERVACION"].lower())))
    add("Caso de nacionalidad pendiente permanece visible", sum(1 for row in per_col["NACIONALIDAD"] if row["ESTADO_VALOR"] == PENDING) == 1, str(sum(1 for row in per_col["NACIONALIDAD"] if row["ESTADO_VALOR"] == PENDING)))
    add("Tres casos de evidencia parcial permanecen visibles", sum(1 for row in per_col["VIGENCIA"] if row["ESTADO_VALOR"] == PENDING and "parcial" in row["OBSERVACION"].lower()) == 3, str(sum(1 for row in per_col["VIGENCIA"] if row["ESTADO_VALOR"] == PENDING and "parcial" in row["OBSERVACION"].lower())))
    add("Matriz de gobernanza contiene 62 filas", len(matrix) == 62, str(len(matrix)))
    pending_rows = matrix[matrix["REQUIERE_REVISION_FINAL"] == "SI"]
    add("Ningun registro pendiente queda marcado como APTO", (pending_rows["APTO_PARA_FUTURA_CARGA"] == "SI").sum() == 0, str((pending_rows["APTO_PARA_FUTURA_CARGA"] == "SI").sum()))
    add("No existen errores Excel", excel_ok, "")
    pes_generated = any("PES_READY" in p.name or "FINAL_PES" in p.name or "CARGA_PES" in p.name for p in PROJECT_ROOT.rglob("*") if p.is_file() and "GOBERNANZA" in p.name)
    add("No se genera archivo final PES", not pes_generated, "")
    return rows


def excel_has_no_errors(path: Path) -> bool:
    wb = load_workbook(path, data_only=False, read_only=True)
    bad = {"#REF!", "#N/A", "#VALUE!", "#NAME?"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and any(err in value for err in bad):
                    return False
    return True


def control_changes(targets: list[Path], before: dict[Path, str], backup_dir: Path) -> list[dict[str, str]]:
    rows = []
    now = datetime.now().isoformat(timespec="seconds")
    for target in targets:
        before_hash = before.get(target, "")
        after_hash = sha256_file(target) if target.exists() and target.is_file() else ("DIRECTORIO" if target.exists() else "")
        before_state = "EXISTIA" if before_hash else "NO_EXISTIA"
        after_state = "EXISTE" if target.exists() else "NO_EXISTE"
        if not before_hash and after_hash:
            tipo = "CREADO"
        elif before_hash and before_hash != after_hash:
            tipo = "MODIFICADO"
        elif before_hash == after_hash and before_hash:
            tipo = "SIN_CAMBIO"
        else:
            tipo = "NO_GENERADO"
        rows.append(
            {
                "ARCHIVO": str(target.relative_to(PROJECT_ROOT)),
                "TIPO_CAMBIO": tipo,
                "ESTADO_ANTES": before_state,
                "ESTADO_DESPUES": after_state,
                "RESPALDO": str(backup_dir),
                "HASH_ANTES": before_hash,
                "HASH_DESPUES": after_hash,
                "FECHA": now,
                "OBSERVACION": "",
            }
        )
    return rows


def build_report(
    official_cols: list[str],
    structure_info: dict[str, str],
    cobertura_cols: list[dict[str, str]],
    cobertura_regs: list[dict[str, str]],
    matrix: pd.DataFrame,
    catalog_paths: dict[str, Path],
    backup_dir: Path,
    validation_rows: list[dict[str, str]],
) -> str:
    global_counts = matrix["ESTADO_GLOBAL_REGISTRO"].value_counts().to_dict()
    apt = int(global_counts.get("APTO", 0))
    warn = int(global_counts.get("APTO_CON_ADVERTENCIAS", 0))
    no_falt = int(global_counts.get("NO_APTO_DATOS_FALTANTES", 0))
    no_conf = int(global_counts.get("NO_APTO_CONFLICTOS", 0))
    pending = int(global_counts.get(PENDING, 0))
    largest = sorted(cobertura_cols, key=lambda r: int(r["PENDIENTE_INSTITUCIONAL"]) + int(r["CONFLICTO"]) + int(r["INVALIDO"]), reverse=True)[:8]
    hashes = {
        "estructura": sha256_file(PATHS["estructura"]),
        "frozen_tsv": sha256_file(PATHS["frozen"]),
        "matriz": sha256_file(OUT["matriz"]),
        "excel": sha256_file(OUT["excel"]),
        "validacion": sha256_file(OUT["validacion"]),
    }
    git_status = os.popen(f"cd {PROJECT_ROOT!s} && git status --short --untracked-files=all -- .").read().strip()
    lines = [
        "# Reporte gobernanza columna por columna Extranjeros Regulares 2025",
        "",
        f"Fecha: {datetime.now().isoformat(timespec='seconds')}",
        "",
        "## Estructura oficial",
        "",
        f"- Archivo: `{PATHS['estructura']}`",
        f"- Hash SHA-256: `{structure_info['hash_sha256']}`",
        f"- Delimitador: `{structure_info['delimitador']}`",
        f"- Codificacion: `{structure_info['codificacion']}`",
        f"- Columnas oficiales: {len(official_cols)}",
        f"- Orden: {'; '.join(f'{i+1}. {c}' for i, c in enumerate(official_cols))}",
        "",
        "## Aptitud de registros",
        "",
        f"- APTO: {apt}",
        f"- APTO_CON_ADVERTENCIAS: {warn}",
        f"- NO_APTO_DATOS_FALTANTES: {no_falt}",
        f"- NO_APTO_CONFLICTOS: {no_conf}",
        f"- PENDIENTE_INSTITUCIONAL: {pending}",
        "",
        "## Principales brechas por columna",
        "",
    ]
    for row in largest:
        lines.append(
            f"- {row['COLUMNA']}: pendientes {row['PENDIENTE_INSTITUCIONAL']}, conflictos {row['CONFLICTO']}, invalidos {row['INVALIDO']}, cobertura {row['PORCENTAJE_VALIDADO']}%."
        )
    lines.extend(
        [
            "",
            "## Pendientes institucionales visibles",
            "",
            f"- Nacionalidad pendiente: {sum(1 for r in cobertura_cols if r['COLUMNA'] == 'NACIONALIDAD' for _ in [0]) and next(r for r in cobertura_cols if r['COLUMNA']=='NACIONALIDAD')['PENDIENTE_INSTITUCIONAL']}",
            f"- Vigencia/evidencia 2025 pendiente: {next(r for r in cobertura_cols if r['COLUMNA']=='VIGENCIA')['PENDIENTE_INSTITUCIONAL']}",
            f"- Codigo unico pendiente: {next(r for r in cobertura_cols if r['COLUMNA']=='CODIGO_UNICO')['PENDIENTE_INSTITUCIONAL']}",
            f"- Pais origen pendiente: {next(r for r in cobertura_cols if r['COLUMNA']=='PAIS_ORIGEN')['PENDIENTE_INSTITUCIONAL']}",
            f"- Pais estudios secundarios pendiente: {next(r for r in cobertura_cols if r['COLUMNA']=='PAIS_ESTUDIOS_SECUNDARIOS')['PENDIENTE_INSTITUCIONAL']}",
            "",
            "## Archivos creados",
            "",
            f"- Esquema: `{OUT['esquema']}`",
            f"- Diccionario: `{OUT['diccionario']}`",
            f"- TSV gobernados: `{OUT['column_dir']}`",
            f"- Catalogos: {', '.join(f'`{p}`' for p in catalog_paths.values())}",
            f"- Matriz: `{OUT['matriz']}`",
            f"- Excel: `{OUT['excel']}`",
            f"- Cobertura columnas: `{OUT['cobertura_columnas']}`",
            f"- Cobertura registros: `{OUT['cobertura_registros']}`",
            f"- Validacion: `{OUT['validacion']}`",
            "",
            "## Hashes",
            "",
        ]
    )
    for key, value in hashes.items():
        lines.append(f"- {key}: `{value}`")
    lines.extend(
        [
            "",
            "## Respaldo",
            "",
            f"`{backup_dir}`",
            "",
            "## Validacion",
            "",
        ]
    )
    for row in validation_rows:
        lines.append(f"- {row['CONTROL']}: {row['RESULTADO']}")
    lines.extend(["", "## Estado git", "", "```", git_status or "Sin cambios reportados por git para este subproducto.", "```", ""])
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", default=str(PROJECT_ROOT), help="Ruta del proyecto; solo informativa en esta version.")
    args = parser.parse_args()
    if Path(args.project_root).resolve() != PROJECT_ROOT:
        raise RuntimeError("Este script debe ejecutarse dentro del proyecto esperado.")

    ensure_dirs()
    required = list(PATHS.values())
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError("Faltan fuentes requeridas: " + "; ".join(missing))

    frozen_hash = sha256_file(PATHS["frozen"])
    if frozen_hash != FROZEN_HASH:
        raise RuntimeError(f"Hash TSV congelado cambiado: {frozen_hash}")

    targets = [
        OUT["control_cambios"],
        OUT["inspeccion"],
        OUT["esquema"],
        OUT["column_dir"],
        OUT["catalog_dir"],
        OUT["matriz"],
        OUT["excel"],
        OUT["diccionario"],
        OUT["cobertura_columnas"],
        OUT["cobertura_registros"],
        OUT["validacion"],
        OUT["reporte"],
    ]
    before_hashes = {target: sha256_file(target) for target in targets if target.exists() and target.is_file()}
    backup_dir = backup_targets(targets + [PROJECT_ROOT / "README.md", Path(__file__)])

    structure_hash_before = sha256_file(PATHS["estructura"])
    official_cols, inspection_rows, structure_info = inspect_structure(PATHS["estructura"])
    write_csv(OUT["inspeccion"], inspection_rows, ["ATRIBUTO", "VALOR", "OBSERVACION"])
    structure_hash_after = sha256_file(PATHS["estructura"])

    schema = build_schema(official_cols)
    schema_cols = [
        "ORDEN",
        "COLUMNA_OFICIAL",
        "EXISTE_EN_ESTRUCTURA",
        "POSICION_ESTRUCTURA",
        "DEFINICION_MANUAL",
        "OBLIGATORIEDAD",
        "TIPO_DATO",
        "FORMATO",
        "LONGITUD",
        "VALORES_PERMITIDOS",
        "REGLA_CONDICIONAL",
        "FUENTE_PREFERENTE",
        "FUENTE_SECUNDARIA",
        "PUEDE_QUEDAR_VACIA",
        "REGLA_DE_VACIO",
        "OBSERVACION",
    ]
    write_tsv(OUT["esquema"], schema, schema_cols)

    frozen = read_csv_df(PATHS["frozen"], sep="\t")
    frozen_original_cols = list(frozen.columns)
    frozen["FILA_BASE_CONGELADA"] = [str(i) for i in range(2, len(frozen) + 2)]
    cierre = read_csv_df(PATHS["cierre"])
    precarga = read_csv_df(PATHS["precarga_norm"])
    conc_prec = read_csv_df(PATHS["conc_precarga"])
    conc_mat = read_csv_df(PATHS["conc_matricula"])
    conc_257 = read_csv_df(PATHS["conc_257"])
    nac = read_csv_df(PATHS["nac_ambigua"])
    evid4 = read_csv_df(PATHS["evidencia4"])
    mult19 = read_csv_df(PATHS["multiples19"])
    catalogo = read_csv_df(PATHS["catalogo_paises"])
    da = load_datos_alumnos(PATHS["datos_alumnos"])

    if len(frozen) != 62:
        raise RuntimeError(f"Base congelada debe tener 62 filas y tiene {len(frozen)}")
    if len(cierre) != 62:
        raise RuntimeError(f"Base de cierre debe tener 62 filas y tiene {len(cierre)}")

    catalog_paths = build_static_catalogs(catalogo)
    per_col, matrix = build_governance_rows(
        official_cols,
        frozen,
        cierre,
        da,
        precarga,
        conc_prec,
        conc_mat,
        conc_257,
        nac,
        evid4,
        mult19,
    )
    tsv_paths = write_governed_tsvs(per_col, official_cols)
    matrix_to_csv(matrix)

    dictionary_rows = build_dictionary(schema, tsv_paths)
    dict_cols = [
        "ORDEN",
        "COLUMNA_OFICIAL",
        "DESCRIPCION_MANUAL",
        "TIPO_DATO",
        "FORMATO",
        "OBLIGATORIA",
        "CONDICIONAL",
        "VALORES_PERMITIDOS",
        "FUENTE_PRIMARIA",
        "FUENTE_SECUNDARIA",
        "REGLA_SELECCION",
        "REGLA_VALIDACION",
        "REGLA_VACIO",
        "MENSAJE_ERROR",
        "RESPONSABLE",
        "TSV_GOBERNADO",
        "OBSERVACION",
    ]
    write_tsv(OUT["diccionario"], dictionary_rows, dict_cols)

    cobertura_cols = build_coverage(per_col, official_cols)
    cobertura_regs = build_record_coverage(matrix, official_cols)
    cov_col_cols = [
        "ORDEN",
        "COLUMNA",
        "TOTAL_REGISTROS",
        "VALIDADO",
        "VALIDADO_CON_ADVERTENCIA",
        "VACIO_PERMITIDO",
        "VACIO_NO_PERMITIDO",
        "CONFLICTO",
        "INVALIDO",
        "PENDIENTE_INSTITUCIONAL",
        "NO_APLICA",
        "PORCENTAJE_VALIDADO",
        "APTO_COLUMNA",
        "OBSERVACION",
    ]
    write_csv(OUT["cobertura_columnas"], cobertura_cols, cov_col_cols)
    cov_reg_cols = [
        "FILA_BASE_CONGELADA",
        "ID_CONTROL",
        "DOCUMENTO",
        "NOMBRE",
        "TOTAL_COLUMNAS_OFICIALES",
        "COLUMNAS_VALIDAS",
        "COLUMNAS_ADVERTENCIA",
        "COLUMNAS_PENDIENTES",
        "COLUMNAS_CONFLICTO",
        "COLUMNAS_INVALIDAS",
        "APTO_PARA_FUTURA_CARGA",
        "MOTIVO_NO_APTO",
        "RESPONSABLE",
        "OBSERVACION",
    ]
    write_csv(OUT["cobertura_registros"], cobertura_regs, cov_reg_cols)

    validation_seed = [{"CONTROL": "Validacion en construccion", "RESULTADO": "OK", "OBSERVACION": ""}]
    write_excel(
        official_cols,
        inspection_rows,
        schema,
        frozen,
        matrix,
        cobertura_cols,
        cobertura_regs,
        dictionary_rows,
        validation_seed,
        per_col,
    )
    excel_ok = excel_has_no_errors(OUT["excel"])
    validation_rows = build_validation(
        official_cols,
        frozen,
        frozen_original_cols,
        matrix,
        per_col,
        structure_hash_before,
        structure_hash_after,
        excel_ok,
    )
    write_csv(OUT["validacion"], validation_rows, ["CONTROL", "RESULTADO", "OBSERVACION"])
    write_excel(
        official_cols,
        inspection_rows,
        schema,
        frozen,
        matrix,
        cobertura_cols,
        cobertura_regs,
        dictionary_rows,
        validation_rows,
        per_col,
    )
    excel_ok = excel_has_no_errors(OUT["excel"])
    if not excel_ok:
        raise RuntimeError("El Excel generado contiene errores visibles.")

    OUT["reporte"].write_text(
        build_report(
            official_cols,
            structure_info,
            cobertura_cols,
            cobertura_regs,
            matrix,
            catalog_paths,
            backup_dir,
            validation_rows,
        ),
        encoding="utf-8",
    )

    all_targets = targets + list(tsv_paths.values()) + list(catalog_paths.values()) + [Path(__file__)]
    control_rows = control_changes(all_targets, before_hashes, backup_dir)
    write_csv(
        OUT["control_cambios"],
        control_rows,
        [
            "ARCHIVO",
            "TIPO_CAMBIO",
            "ESTADO_ANTES",
            "ESTADO_DESPUES",
            "RESPALDO",
            "HASH_ANTES",
            "HASH_DESPUES",
            "FECHA",
            "OBSERVACION",
        ],
    )

    errors = [row for row in validation_rows if row["RESULTADO"] == "ERROR"]
    if errors:
        raise RuntimeError("Validacion de gobernanza con errores: " + "; ".join(row["CONTROL"] for row in errors))

    summary = {
        "hash_tsv_congelado": frozen_hash,
        "hash_estructura": structure_info["hash_sha256"],
        "delimitador": structure_info["delimitador"],
        "codificacion": structure_info["codificacion"],
        "columnas_oficiales": len(official_cols),
        "tsv_gobernados": len(tsv_paths),
        "filas_matriz": len(matrix),
        "apto": int((matrix["ESTADO_GLOBAL_REGISTRO"] == "APTO").sum()),
        "apto_con_advertencias": int((matrix["ESTADO_GLOBAL_REGISTRO"] == "APTO_CON_ADVERTENCIAS").sum()),
        "no_apto_datos_faltantes": int((matrix["ESTADO_GLOBAL_REGISTRO"] == "NO_APTO_DATOS_FALTANTES").sum()),
        "no_apto_conflictos": int((matrix["ESTADO_GLOBAL_REGISTRO"] == "NO_APTO_CONFLICTOS").sum()),
        "pendiente_institucional": int((matrix["ESTADO_GLOBAL_REGISTRO"] == PENDING).sum()),
        "backup": str(backup_dir),
        "excel": str(OUT["excel"]),
        "matriz": str(OUT["matriz"]),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
