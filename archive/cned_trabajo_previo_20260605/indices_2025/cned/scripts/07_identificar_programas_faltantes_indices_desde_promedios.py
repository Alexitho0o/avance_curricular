#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera un artefacto auditable de programas CNED faltantes desde PROMEDIOS/matriz."""

from __future__ import annotations

import json
import re
import subprocess
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


CODE_PATTERN = re.compile(r"^I(?P<ies>\d+)S(?P<sed>\d+)C(?P<car>\d+)J(?P<jor>\d+)V(?P<version>\d+)$")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="F4CCCC")
INFO_FILL = PatternFill("solid", fgColor="D9EAF7")

MATRIX_COLUMNS = [
    "CODIGO_IES_NUM",
    "CODIGO_UNICO",
    "NOMBRE_IES",
    "COD_SEDE",
    "NOMBRE_SEDE",
    "NOMBRE_CARRERA",
    "MODALIDAD",
    "JORNADA",
    "TIPO_PLAN_CARRERA",
    "CARACT_PLAN_ESPECIAL",
    "DURACION_ESTUDIOS",
    "DURACION_TITULACION",
    "DURACION_TOTAL",
    "COMUNA_SEDE",
    "PROVINCIA_SEDE",
    "REGION_SEDE",
    "NIVEL_GLOBAL",
    "NIVEL_CARRERA",
    "VIGENCIA",
]

REFERENCE_COLUMNS = [
    "Sede",
    "Comuna",
    "Región",
    "Orden Geográfico de Region",
    "Cod.Carr.",
    "Carrera Genérica",
    "Nombre Programa",
    "Mencion Especialidad",
    "Horario",
    "Tipo Programa",
    "Área Conocimiento",
    "Tipo Carrera",
    "Ingreso Directo",
    "Año Inicio Actividades",
    "Campus",
    "Duracion (en semestres)",
    "Titulo",
    "Grado Academico",
    "Codigo SIES",
    "Nivel Pregrado/Posgrado",
    "Modalidad Programa",
    "Especial",
    "Facultad",
]

FORM_COLUMNS = [
    "Sede",
    "Tipo de Carrera",
    "Nombre",
    "Mención o Especialidad",
    "Año de inicio de Actividades",
    "Campus",
    "Horario",
    "Estado",
    "Tipo Programa",
    "Detalle del Tipo de Programa (especiales)",
    "Modalidad del Programa",
    "Área del Conocimiento",
    "Sub Área",
    "Carrera Genérica",
    "Título que otorga el programa",
    "Grado Académico que otorga el programa",
    "Dependencia",
    "Régimen",
    "Duración del Programa en Semestres",
    "Ingreso Directo",
    "Ingreso desde un plan Común",
    "Ingreso desde Bachillerato",
    "Otro Tipo de Ingreso",
    "Ingreso otro",
    "Observaciones",
]

TRACE_COLUMNS = [
    "CODIGO_UNICO",
    "CODIGO_IES_NUM",
    "COD_SEDE",
    "COD_CARRERA_DERIVADO",
    "JORNADA",
    "VERSION_DERIVADA",
    "NOMBRE_CARRERA_ORIGINAL",
    "NOMBRE_CARRERA_NORMALIZADO",
    "MODALIDAD_ORIGINAL",
    "TIPO_PLAN_CARRERA_ORIGINAL",
    "CARACT_PLAN_ESPECIAL_ORIGINAL",
    "NIVEL_GLOBAL",
    "NIVEL_CARRERA",
    "VIGENCIA",
    "ESTADO_CRUCE_INDICES",
    "MOTIVO_ESTADO_CRUCE",
    "FUENTE",
    "REVISADO_MANUALMENTE",
    "OBSERVACION_AUDITORIA",
]

REQUIRED_SHEETS = [
    "RESUMEN_EJECUTIVO",
    "OFERTA_MATRIZ_NORMALIZADA",
    "PROGRAMAS_EXISTENTES_NORMALIZADOS",
    "CRUCE_MATRIZ_VS_INDICES",
    "PROGRAMAS_FALTANTES_PARA_CREAR",
    "DUDOSOS_REVISAR_MANUAL",
    "YA_EXISTEN_EN_INDICES",
    "NO_CREAR",
    "DICCIONARIO_CAMPOS",
    "AUDITORIA_COLUMNAS",
    "AUDITORIA_DUPLICADOS",
    "AUDITORIA_CODIGO_UNICO",
    "AUDITORIA_CODIGOS_ESPERADOS",
    "BITACORA_PROCESO",
]

EXPECTED_CODES = [
    ("I162S2C101J4V1", "DIPLOMADO EN SALUD FAMILIAR CON ENFOQUE COMUNITARIO"),
    ("I162S2C102J4V1", "DIPLOMADO SALUD CON FOCO EN MIGRACION"),
    ("I162S2C103J4V1", "DIPLOMADO ABORDAJE INTEGRAL DE LA VIOLENCIA DE GENERO EN ATENCION PRIMARIA DE SALUD"),
    ("I162S2C104J4V1", "DIPLOMADO EN INNOVACION PARA LA DOCENCIA"),
    ("I162S2C105J4V1", "DIPLOMADO EN HUMANIZACION EN SALUD"),
    ("I162S2C111J4V1", "INGENIERIA EN FINANZAS"),
    ("I162S2C111J4V2", "INGENIERIA EN FINANZAS"),
    ("I162S2C112J4V1", "INGENIERIA EN MARKETING DIGITAL"),
    ("I162S2C112J4V2", "INGENIERIA EN MARKETING DIGITAL"),
    ("I162S2C113J4V1", "INGENIERIA EN RECURSOS HUMANOS"),
    ("I162S2C113J4V2", "INGENIERIA EN RECURSOS HUMANOS"),
    ("I162S2C114J2V1", "TECNICO EN ENFERMERIA E INSTRUMENTACION QUIRURGICA"),
    ("I162S2C115J2V1", "TECNICO EN FARMACIA"),
    ("I162S2C116J4V1", "TECNICO EN FINANZAS"),
    ("I162S2C117J4V1", "TECNICO EN INFRAESTRUCTURA CLOUD"),
    ("I162S2C118J4V1", "TECNICO EN MARKETING DIGITAL"),
    ("I162S2C119J4V1", "TECNICO EN RECURSOS HUMANOS"),
    ("I162S2C124J4V1", "INGENIERIA EN PREVENCION DE RIESGOS"),
    ("I162S2C124J4V2", "INGENIERIA EN PREVENCION DE RIESGOS"),
    ("I162S2C125J4V1", "INGENIERIA EN SEGURIDAD PRIVADA"),
    ("I162S2C125J4V2", "INGENIERIA EN SEGURIDAD PRIVADA"),
    ("I162S2C95J4V1", "DIPLOMADO EN INTELIGENCIA ARTIFICIAL"),
    ("I162S2C96J4V1", "DIPLOMADO EN GOBERNANZA DE DATOS"),
    ("I162S2C97J4V1", "DIPLOMADO EN SUPPLY CHAIN MANAGEMENT Y MINERIA DE REQUERIMIENTOS"),
    ("I162S2C98J4V1", "DIPLOMADO EN HABILIDADES DIRECTIVAS PARA PROFESIONALES STEM"),
    ("I162S2C99J4V1", "DIPLOMADO EN GESTION PUBLICA LOCAL"),
    ("I162S3C114J2V1", "TECNICO EN ENFERMERIA E INSTRUMENTACION QUIRURGICA"),
    ("I162S3C115J2V1", "TECNICO EN FARMACIA"),
    ("I162S3C91J2V1", "TECNICO EN ENFERMERIA"),
]

RULES_CROSS = (
    "prioridad: CODIGO_UNICO exacto > auxiliar fuerte por nombre+sede+estructura > faltante; "
    "sin equivalencias forzadas entre codigos de distinta naturaleza; campos sin evidencia quedan REVISAR"
)


@dataclass(frozen=True)
class Paths:
    repo: Path
    matrix_path: Path
    existing_path: Path
    manual_tsv_path: Path
    resultados: Path
    excel_out: Path
    csv_out: Path
    md_out: Path
    json_out: Path
    script_path: str
    timestamp: str


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.lower() == "nan":
        return ""
    return text.strip()


def clean_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for column in out.columns:
        out[column] = out[column].map(clean_cell)
    return out


def normalize_text(value: Any) -> str:
    text = clean_cell(value)
    if not text:
        return ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text)
    text = text.upper()
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_program_name(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        return ""
    replacements = {
        "TECNICO NIVEL SUPERIOR": "TNS",
        "TECNICO NIVEL SUP": "TNS",
        "TECNICO NIV SUP": "TNS",
        "ING EJEC": "INGENIERIA EJECUCION",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\bP E\b", " ", text)
    text = re.sub(r"\bPE\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_sede_key(value: Any) -> str:
    original = clean_cell(value)
    if not original:
        return ""
    match = re.search(r"\((.*?)\)", original)
    if match:
        return normalize_text(match.group(1))
    normalized = normalize_text(original)
    normalized = re.sub(r"^CASA CENTRAL\s+", "", normalized)
    normalized = re.sub(r"^SEDE\s+", "", normalized)
    return normalized.strip()


def pretty_sede(value: Any) -> str:
    key = normalize_sede_key(value)
    if key == "SANTIAGO":
        return "Santiago"
    if key == "CONCEPCION":
        return "Concepcion"
    original = clean_cell(value)
    if original:
        return original.title()
    return "REVISAR"


def parse_code(code: str) -> dict[str, str]:
    match = CODE_PATTERN.fullmatch(clean_cell(code))
    if not match:
        return {"CODIGO_IES_DERIVADO": "", "COD_SEDE_DERIVADO": "", "COD_CARRERA_DERIVADO": "", "JORNADA_DERIVADA": "", "VERSION_DERIVADA": "", "FORMATO_CODIGO_UNICO": "REVISAR"}
    data = match.groupdict()
    return {
        "CODIGO_IES_DERIVADO": data["ies"],
        "COD_SEDE_DERIVADO": data["sed"],
        "COD_CARRERA_DERIVADO": data["car"],
        "JORNADA_DERIVADA": data["jor"],
        "VERSION_DERIVADA": data["version"],
        "FORMATO_CODIGO_UNICO": "OK",
    }


def derive_code_df(series: pd.Series) -> pd.DataFrame:
    return pd.DataFrame([parse_code(clean_cell(value)) for value in series])


def locate_repo() -> Path:
    candidates: list[Path] = [Path.cwd()]
    file_name = globals().get("__file__")
    if file_name:
        try:
            candidates.append(Path(file_name).resolve().parents[3])
        except Exception:
            pass
    for candidate in candidates:
        if (candidate / "indices_2025" / "cned").exists():
            return candidate
    raise FileNotFoundError("No fue posible ubicar el repo avance_curricular.")


def build_paths() -> Paths:
    repo = locate_repo()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    script_name = globals().get("__file__")
    script_path = str(Path(script_name).resolve()) if script_name else "stdin"
    resultados = repo / "indices_2025" / "cned" / "resultados"
    base_name = f"CNED_PROGRAMAS_FALTANTES_CREAR_INDICES_DESDE_PROMEDIOS_{ts}"
    return Paths(
        repo=repo,
        matrix_path=repo / "input" / "PROMEDIOSDEALUMNOS_7804.xlsx",
        existing_path=repo / "indices_2025" / "cned" / "data" / "listado_referencia_cned.tsv",
        manual_tsv_path=repo / "indices_2025" / "cned" / "data" / "DATOS_VALIDACION_CNED_TSV_NUEVO_MANUAL.tsv",
        resultados=resultados,
        excel_out=resultados / f"{base_name}.xlsx",
        csv_out=resultados / f"{base_name}.csv",
        md_out=resultados / f"{base_name}.md",
        json_out=resultados / f"{base_name}.json",
        script_path=script_path,
        timestamp=ts,
    )


def git_output(repo: Path, *args: str) -> str:
    result = subprocess.run(["git", *args], cwd=repo, capture_output=True, text=True, check=False)
    if result.returncode != 0:
        return clean_cell(result.stderr)
    return clean_cell(result.stdout)


def read_matrix(paths: Paths) -> tuple[pd.DataFrame, list[dict[str, Any]], list[str]]:
    if not paths.matrix_path.exists():
        raise FileNotFoundError(f"No existe {paths.matrix_path}")
    workbook = load_workbook(paths.matrix_path, read_only=True, data_only=False)
    sheet_meta: list[dict[str, Any]] = []
    for name in workbook.sheetnames:
        ws = workbook[name]
        sheet_meta.append({"sheet": name, "rows": max(ws.max_row - 1, 0), "cols": ws.max_column})
    if "matriz" not in workbook.sheetnames:
        raise ValueError(f"No existe hoja matriz en {paths.matrix_path}; hojas encontradas: {workbook.sheetnames}")
    df = clean_df(pd.read_excel(paths.matrix_path, sheet_name="matriz", dtype=str))
    missing = [column for column in MATRIX_COLUMNS if column not in df.columns]
    if missing:
        raise ValueError(f"La hoja matriz no contiene columnas requeridas: {missing}")
    return df[MATRIX_COLUMNS].copy(), sheet_meta, workbook.sheetnames


def read_existing(paths: Paths) -> pd.DataFrame:
    if not paths.existing_path.exists():
        raise FileNotFoundError(f"No existe {paths.existing_path}")
    df = clean_df(pd.read_csv(paths.existing_path, sep="\t", dtype=str, keep_default_na=False))
    missing = [column for column in REFERENCE_COLUMNS if column not in df.columns]
    if missing:
        raise ValueError(f"El listado existente no contiene columnas requeridas: {missing}")
    return df[REFERENCE_COLUMNS].copy()


def build_offer_matrix_normalized(df: pd.DataFrame, paths: Paths) -> pd.DataFrame:
    out = df.copy()
    deriv = derive_code_df(out["CODIGO_UNICO"])
    out = pd.concat([out.reset_index(drop=True), deriv.reset_index(drop=True)], axis=1)
    out["VALIDACION_IES_CODIGO_UNICO"] = out.apply(
        lambda row: "OK" if row["FORMATO_CODIGO_UNICO"] == "OK" and row["CODIGO_IES_NUM"] == row["CODIGO_IES_DERIVADO"] else "REVISAR",
        axis=1,
    )
    out["VALIDACION_SEDE_CODIGO_UNICO"] = out.apply(
        lambda row: "OK" if row["FORMATO_CODIGO_UNICO"] == "OK" and row["COD_SEDE"] == row["COD_SEDE_DERIVADO"] else "REVISAR",
        axis=1,
    )
    out["VALIDACION_JORNADA_CODIGO_UNICO"] = out.apply(
        lambda row: "OK" if row["FORMATO_CODIGO_UNICO"] == "OK" and row["JORNADA"] == row["JORNADA_DERIVADA"] else "REVISAR",
        axis=1,
    )
    out["NOMBRE_CARRERA_NORMALIZADO"] = out["NOMBRE_CARRERA"].map(normalize_program_name)
    out["NOMBRE_CARRERA_NORMALIZADO_BASE"] = out["NOMBRE_CARRERA"].map(normalize_text)
    out["NOMBRE_SEDE_NORMALIZADO"] = out["NOMBRE_SEDE"].map(normalize_text)
    out["SEDE_CRUCE_NORMALIZADA"] = out["NOMBRE_SEDE"].map(normalize_sede_key)
    out["CLAVE_CODIGO_UNICO"] = out["CODIGO_UNICO"]
    out["CLAVE_AUXILIAR_NOMBRE_SEDE_JORNADA"] = (
        out["NOMBRE_CARRERA_NORMALIZADO"]
        + "|"
        + out["SEDE_CRUCE_NORMALIZADA"]
        + "|"
        + out["JORNADA"]
    )
    out["CLAVE_AUXILIAR_ESTRUCTURAL"] = (
        out["NOMBRE_CARRERA_NORMALIZADO"]
        + "|"
        + out["SEDE_CRUCE_NORMALIZADA"]
        + "|"
        + out["JORNADA"]
        + "|"
        + out["MODALIDAD"]
        + "|"
        + out["TIPO_PLAN_CARRERA"]
        + "|"
        + out["DURACION_ESTUDIOS"]
        + "|"
        + out["NIVEL_GLOBAL"]
        + "|"
        + out["NIVEL_CARRERA"]
    )
    out["FUENTE"] = f"{paths.matrix_path.name} / hoja matriz"
    return out


def infer_reference_nivel_carrera(tipo_carrera: str) -> str:
    normalized = normalize_text(tipo_carrera)
    if normalized == "TECNICO NIVEL SUPERIOR":
        return "1"
    if normalized == "PROFESIONAL":
        return "2"
    if normalized == "DIPLOMADO":
        return "5"
    return ""


def infer_reference_jornada(horario: str) -> str:
    normalized = normalize_text(horario)
    mapping = {"DIURNO": "1", "VESPERTINO": "2", "OTRO": "4"}
    return mapping.get(normalized, "")


def infer_reference_modalidad(modalidad: str) -> str:
    normalized = normalize_text(modalidad)
    mapping = {"PRESENCIAL": "1", "NO PRESENCIAL": "3"}
    return mapping.get(normalized, "")


def infer_reference_tipo_plan(tipo_programa: str) -> str:
    normalized = normalize_text(tipo_programa)
    mapping = {"PROGRAMA REGULAR": "1", "PROGRAMA ESPECIAL": "3"}
    return mapping.get(normalized, "")


def build_existing_normalized(df: pd.DataFrame, paths: Paths) -> pd.DataFrame:
    out = df.copy()
    out["CODIGO_UNICO"] = out["Codigo SIES"]
    deriv = derive_code_df(out["CODIGO_UNICO"])
    out = pd.concat([out.reset_index(drop=True), deriv.reset_index(drop=True)], axis=1)
    out["CODIGO_UNICO_DUPLICADO"] = out["CODIGO_UNICO"].where(out["CODIGO_UNICO"] != "", "").duplicated(keep=False).map(lambda flag: "SI" if flag else "NO")
    out.loc[out["CODIGO_UNICO"].eq(""), "CODIGO_UNICO_DUPLICADO"] = "NO"
    out["NOMBRE_PROGRAMA_NORMALIZADO"] = out["Nombre Programa"].map(normalize_program_name)
    out["NOMBRE_PROGRAMA_NORMALIZADO_BASE"] = out["Nombre Programa"].map(normalize_text)
    out["SEDE_CRUCE_NORMALIZADA"] = out["Sede"].map(normalize_sede_key)
    out["HORARIO_NORMALIZADO"] = out["Horario"].map(normalize_text)
    out["MODALIDAD_NORMALIZADA"] = out["Modalidad Programa"].map(normalize_text)
    out["TIPO_PROGRAMA_NORMALIZADO"] = out["Tipo Programa"].map(normalize_text)
    out["TIPO_CARRERA_NORMALIZADA"] = out["Tipo Carrera"].map(normalize_text)
    out["NIVEL_CARRERA_ESTIMADO"] = out["Tipo Carrera"].map(infer_reference_nivel_carrera)
    out["JORNADA_ESTIMADA"] = out["Horario"].map(infer_reference_jornada)
    out["MODALIDAD_ESTIMADA"] = out["Modalidad Programa"].map(infer_reference_modalidad)
    out["TIPO_PLAN_ESTIMADO"] = out["Tipo Programa"].map(infer_reference_tipo_plan)
    out["CLAVE_CODIGO_UNICO"] = out["CODIGO_UNICO"]
    out["CLAVE_AUXILIAR_NOMBRE_SEDE_JORNADA"] = (
        out["NOMBRE_PROGRAMA_NORMALIZADO"]
        + "|"
        + out["SEDE_CRUCE_NORMALIZADA"]
        + "|"
        + out["JORNADA_ESTIMADA"]
    )
    out["CLAVE_AUXILIAR_ESTRUCTURAL"] = (
        out["NOMBRE_PROGRAMA_NORMALIZADO"]
        + "|"
        + out["SEDE_CRUCE_NORMALIZADA"]
        + "|"
        + out["JORNADA_ESTIMADA"]
        + "|"
        + out["MODALIDAD_ESTIMADA"]
        + "|"
        + out["TIPO_PLAN_ESTIMADO"]
        + "|"
        + out["Duracion (en semestres)"]
        + "|"
        + out["NIVEL_CARRERA_ESTIMADO"]
    )
    out["FUENTE_EXISTENTES"] = str(paths.existing_path)
    return out


def majority_label(df: pd.DataFrame, code_column: str, label_column: str) -> tuple[dict[str, str], list[dict[str, str]], list[str]]:
    mapping: dict[str, str] = {}
    rows: list[dict[str, str]] = []
    warnings: list[str] = []
    for code, group in df.groupby(code_column):
        if clean_cell(code) == "":
            continue
        counts = group[label_column].map(clean_cell).value_counts(dropna=False)
        label = clean_cell(counts.index[0])
        mapping[str(code)] = label
        ambiguous = "SI" if len(counts) > 1 else "NO"
        rows.append(
            {
                "SECCION": f"DICCIONARIO_{code_column}_{label_column}",
                "CODIGO": str(code),
                "VALOR_SELECCIONADO": label,
                "CONTEOS_OBSERVADOS": json.dumps({str(key): int(value) for key, value in counts.items()}, ensure_ascii=False),
                "AMBIGUO": ambiguous,
                "FUENTE": "solape exacto CODIGO_UNICO matriz vs existentes",
                "OBSERVACION": "mayoria empirica" if ambiguous == "SI" else "mapeo unico",
            }
        )
        if ambiguous == "SI":
            warnings.append(f"Diccionario ambiguo {code_column}->{label_column} para codigo {code}: {dict(counts)}")
    return mapping, rows, warnings


def build_dictionaries(offer_df: pd.DataFrame, existing_df: pd.DataFrame) -> tuple[dict[str, dict[str, str]], pd.DataFrame, list[str]]:
    exact = offer_df.merge(existing_df, on="CODIGO_UNICO", how="inner")
    warnings: list[str] = []
    dict_rows: list[dict[str, str]] = []

    jornada_map = {"1": "Diurno", "2": "Vespertino", "4": "Otro"}
    for code, label in jornada_map.items():
        dict_rows.append(
            {
                "SECCION": "DICCIONARIO_JORNADA",
                "CODIGO": code,
                "VALOR_SELECCIONADO": label,
                "CONTEOS_OBSERVADOS": "solape exacto sin ambiguedad",
                "AMBIGUO": "NO",
                "FUENTE": "solape exacto CODIGO_UNICO matriz vs existentes",
                "OBSERVACION": "jornada 3 no observada en existentes exactos",
            }
        )
    modalidad_map, modalidad_rows, modalidad_warnings = majority_label(exact, "MODALIDAD", "Modalidad Programa")
    tipo_plan_map, tipo_plan_rows, tipo_plan_warnings = majority_label(exact, "TIPO_PLAN_CARRERA", "Tipo Programa")
    dict_rows.extend(modalidad_rows)
    dict_rows.extend(tipo_plan_rows)
    warnings.extend(modalidad_warnings)
    warnings.extend(tipo_plan_warnings)

    nivel_carrera_map = {"1": "Tecnico Nivel Superior", "2": "Profesional", "5": "Diplomado"}
    for code, label in nivel_carrera_map.items():
        dict_rows.append(
            {
                "SECCION": "DICCIONARIO_NIVEL_CARRERA",
                "CODIGO": code,
                "VALOR_SELECCIONADO": label,
                "CONTEOS_OBSERVADOS": "regla proyecto/matriz",
                "AMBIGUO": "NO",
                "FUENTE": "matriz + nivel carrera",
                "OBSERVACION": "nivel 5 se trata como Diplomado porque todos los casos observados contienen DIPLOMADO",
            }
        )

    campus_map: dict[str, str] = {}
    for sede, group in existing_df.groupby("SEDE_CRUCE_NORMALIZADA"):
        values = sorted({clean_cell(value) for value in group["Campus"] if clean_cell(value)})
        if len(values) == 1:
            campus_map[sede] = values[0]
            dict_rows.append(
                {
                    "SECCION": "DICCIONARIO_CAMPUS_POR_SEDE",
                    "CODIGO": sede,
                    "VALOR_SELECCIONADO": values[0],
                    "CONTEOS_OBSERVADOS": "mapeo unico por sede en existentes",
                    "AMBIGUO": "NO",
                    "FUENTE": str(existing_df["FUENTE_EXISTENTES"].iloc[0]),
                    "OBSERVACION": "si la sede no tiene campus unico, se deja REVISAR",
                }
            )

    column_map_rows = [
        {"SECCION": "MAPEO_MATRIZ", "CAMPO_ORIGINAL": column, "CAMPO_NORMALIZADO": column, "CODIGO": "", "VALOR_SELECCIONADO": "", "CONTEOS_OBSERVADOS": "", "AMBIGUO": "", "FUENTE": f"{offer_df['FUENTE'].iloc[0]}", "OBSERVACION": "columna canonicamente usada"}
        for column in MATRIX_COLUMNS
    ]
    column_map_rows.extend(
        {
            "SECCION": "MAPEO_EXISTENTES",
            "CAMPO_ORIGINAL": original,
            "CAMPO_NORMALIZADO": normalized,
            "CODIGO": "",
            "VALOR_SELECCIONADO": "",
            "CONTEOS_OBSERVADOS": "",
            "AMBIGUO": "",
            "FUENTE": str(existing_df["FUENTE_EXISTENTES"].iloc[0]),
            "OBSERVACION": "campo homologado para cruce",
        }
        for original, normalized in [
            ("Sede", "NOMBRE_SEDE"),
            ("Comuna", "COMUNA_SEDE"),
            ("Región", "REGION_SEDE"),
            ("Cod.Carr.", "COD_CARRERA_REFERENCIA"),
            ("Nombre Programa", "NOMBRE_CARRERA"),
            ("Horario", "HORARIO"),
            ("Tipo Programa", "TIPO_PROGRAMA"),
            ("Tipo Carrera", "TIPO_CARRERA"),
            ("Campus", "CAMPUS"),
            ("Codigo SIES", "CODIGO_UNICO"),
        ]
    )
    dict_rows.extend(column_map_rows)

    dictionary_df = pd.DataFrame(dict_rows).fillna("")
    return {
        "jornada": jornada_map,
        "modalidad": modalidad_map,
        "tipo_plan": tipo_plan_map,
        "nivel_carrera": nivel_carrera_map,
        "campus_por_sede": campus_map,
    }, dictionary_df, warnings


def tipo_carrera_formulario(row: pd.Series, dictionaries: dict[str, dict[str, str]]) -> str:
    label = dictionaries["nivel_carrera"].get(clean_cell(row["NIVEL_CARRERA"]), "")
    if label:
        return label
    return "REVISAR"


def horario_formulario(jornada: str, dictionaries: dict[str, dict[str, str]]) -> str:
    return dictionaries["jornada"].get(clean_cell(jornada), "REVISAR")


def modalidad_formulario(modalidad: str, dictionaries: dict[str, dict[str, str]]) -> str:
    label = clean_cell(dictionaries["modalidad"].get(clean_cell(modalidad), ""))
    return label if label else "REVISAR"


def tipo_programa_formulario(tipo_plan: str, dictionaries: dict[str, dict[str, str]]) -> str:
    label = clean_cell(dictionaries["tipo_plan"].get(clean_cell(tipo_plan), ""))
    return label if label else "REVISAR"


def build_offer_matrix_classified(df: pd.DataFrame, dictionaries: dict[str, dict[str, str]]) -> pd.DataFrame:
    out = df.copy()
    out["ES_VIGENTE"] = out["VIGENCIA"].map(lambda value: "SI" if clean_cell(value) == "1" else "NO")
    out["ES_PREGRADO"] = out["NIVEL_GLOBAL"].map(lambda value: "SI" if clean_cell(value) == "1" else "NO")
    out["ES_POSGRADO_POSTITULO_DIPLOMADO"] = out.apply(
        lambda row: "SI" if clean_cell(row["NIVEL_GLOBAL"]) == "3" or clean_cell(row["NIVEL_CARRERA"]) == "5" else "NO",
        axis=1,
    )
    out["ES_DIPLOMADO"] = out.apply(
        lambda row: "SI" if clean_cell(row["NIVEL_CARRERA"]) == "5" or "DIPLOMADO" in row["NOMBRE_CARRERA_NORMALIZADO"] else "NO",
        axis=1,
    )
    out["ES_TECNICO"] = out["NIVEL_CARRERA"].map(lambda value: "SI" if clean_cell(value) == "1" else "NO")
    out["ES_PROFESIONAL"] = out["NIVEL_CARRERA"].map(lambda value: "SI" if clean_cell(value) == "2" else "NO")
    out["ES_PROGRAMA_ESPECIAL"] = out["TIPO_PLAN_CARRERA"].map(lambda value: "SI" if clean_cell(value) == "3" else "NO")
    out["ES_PROGRAMA_REGULAR"] = out["TIPO_PLAN_CARRERA"].map(lambda value: "SI" if clean_cell(value) == "1" else "NO")
    out["ALCANCE_INDICES_CREAR_PROGRAMA"] = out["ES_VIGENTE"].map(lambda value: "SI" if value == "SI" else "NO")
    out["MOTIVO_ALCANCE"] = out["ES_VIGENTE"].map(lambda value: "VIGENTE_EN_MATRIZ" if value == "SI" else "NO_VIGENTE")
    out["SEDE_FORMULARIO"] = out["NOMBRE_SEDE"].map(pretty_sede)
    out["TIPO_DE_CARRERA_FORMULARIO"] = out.apply(tipo_carrera_formulario, axis=1, dictionaries=dictionaries)
    out["HORARIO_FORMULARIO"] = out["JORNADA"].map(lambda value: horario_formulario(value, dictionaries))
    out["TIPO_PROGRAMA_FORMULARIO"] = out["TIPO_PLAN_CARRERA"].map(lambda value: tipo_programa_formulario(value, dictionaries))
    out["MODALIDAD_DEL_PROGRAMA_FORMULARIO"] = out["MODALIDAD"].map(lambda value: modalidad_formulario(value, dictionaries))
    out["DETALLE_TIPO_PROGRAMA_ESPECIAL_FORMULARIO"] = out.apply(
        lambda row: clean_cell(row["CARACT_PLAN_ESPECIAL"]) if row["TIPO_PROGRAMA_FORMULARIO"] == "Programa Especial" and clean_cell(row["CARACT_PLAN_ESPECIAL"]) not in {"", "NO APLICA"} else ("REVISAR" if row["TIPO_PROGRAMA_FORMULARIO"] == "Programa Especial" else ""),
        axis=1,
    )
    out["CAMPUS_FORMULARIO"] = out["SEDE_CRUCE_NORMALIZADA"].map(lambda key: dictionaries["campus_por_sede"].get(key, "REVISAR"))
    return out


def auxiliary_candidates(row: pd.Series, existing_df: pd.DataFrame) -> pd.DataFrame:
    subset = existing_df[
        (existing_df["NOMBRE_PROGRAMA_NORMALIZADO"] == row["NOMBRE_CARRERA_NORMALIZADO"])
        & (existing_df["SEDE_CRUCE_NORMALIZADA"] == row["SEDE_CRUCE_NORMALIZADA"])
    ].copy()
    if subset.empty:
        return subset
    subset["AUX_SCORE"] = 0
    subset["AUX_MATCH_FIELDS"] = ""
    checks = [
        ("JORNADA_ESTIMADA", row["JORNADA"], "jornada"),
        ("MODALIDAD_ESTIMADA", row["MODALIDAD"], "modalidad"),
        ("TIPO_PLAN_ESTIMADO", row["TIPO_PLAN_CARRERA"], "tipo_plan"),
        ("Duracion (en semestres)", row["DURACION_ESTUDIOS"], "duracion"),
        ("NIVEL_CARRERA_ESTIMADO", row["NIVEL_CARRERA"], "nivel_carrera"),
    ]
    for index, candidate in subset.iterrows():
        matched: list[str] = []
        score = 0
        for column, expected, label in checks:
            if clean_cell(expected) and clean_cell(candidate[column]) and clean_cell(candidate[column]) == clean_cell(expected):
                matched.append(label)
                score += 1
        subset.at[index, "AUX_SCORE"] = score
        subset.at[index, "AUX_MATCH_FIELDS"] = ", ".join(matched)
    subset = subset.sort_values(["AUX_SCORE", "CODIGO_UNICO", "Cod.Carr."], ascending=[False, True, True])
    best_score = int(subset["AUX_SCORE"].max())
    if best_score < 3:
        return subset.iloc[0:0]
    return subset[subset["AUX_SCORE"] == best_score].copy()


def build_cross(offer_df: pd.DataFrame, existing_df: pd.DataFrame) -> pd.DataFrame:
    existing_codes = existing_df[existing_df["CODIGO_UNICO"] != ""].copy()
    rows: list[dict[str, str]] = []
    for _, row in offer_df.iterrows():
        base = row.to_dict()
        estado = ""
        motivo = ""
        tipo_match = ""
        confianza = ""
        codigo_match = ""
        observacion = ""
        requiere = "NO"

        if row["ES_VIGENTE"] != "SI":
            estado = "NO_CREAR_POR_NO_VIGENTE"
            motivo = "VIGENCIA_DISTINTA_DE_1"
            tipo_match = "NO_APLICA"
            confianza = "NO_APLICA"
            observacion = "Registro no vigente en matriz."
        elif row["ALCANCE_INDICES_CREAR_PROGRAMA"] != "SI":
            estado = "NO_CREAR_POR_FUERA_DE_ALCANCE"
            motivo = row["MOTIVO_ALCANCE"]
            tipo_match = "NO_APLICA"
            confianza = "NO_APLICA"
            observacion = "Registro fuera del alcance definido."
        else:
            exact = existing_codes[existing_codes["CODIGO_UNICO"] == row["CODIGO_UNICO"]]
            if not exact.empty:
                estado = "YA_EXISTE_EN_INDICES"
                motivo = "COINCIDENCIA_EXACTA_POR_CODIGO_UNICO"
                tipo_match = "CODIGO_UNICO_EXACTO"
                confianza = "ALTA"
                codigo_match = row["CODIGO_UNICO"]
                requiere = "NO"
                observacion = f"Coincidencia exacta encontrada en existentes ({len(exact)} fila(s))."
                if len(exact) > 1:
                    observacion += " Existen duplicados en la base de referencia para este codigo."
            else:
                aux = auxiliary_candidates(row, existing_df)
                if not aux.empty:
                    estado = "DUDOSO_REVISAR_MANUAL"
                    motivo = "COINCIDENCIA_AUXILIAR_FUERTE_SIN_CODIGO_UNICO_EXACTO"
                    tipo_match = "AUXILIAR_NOMBRE_SEDE_ESTRUCTURAL"
                    confianza = "ALTA" if int(aux["AUX_SCORE"].iloc[0]) >= 4 else "MEDIA"
                    codigo_match = " | ".join(code or f"SIN_CODIGO:{clean_cell(code_ref)}" for code, code_ref in zip(aux["CODIGO_UNICO"], aux["Cod.Carr."]))
                    requiere = "NO"
                    observacion = (
                        "Candidato(s) auxiliar(es) por nombre+sede con score "
                        f"{int(aux['AUX_SCORE'].iloc[0])}: { ' || '.join(aux['AUX_MATCH_FIELDS'].tolist()) }"
                    )
                else:
                    estado = "FALTA_CREAR_EN_INDICES"
                    motivo = "SIN_COINCIDENCIA_EXACTA_NI_AUXILIAR_FUERTE"
                    tipo_match = "SIN_MATCH"
                    confianza = "ALTA"
                    requiere = "SI"
                    observacion = "No existe coincidencia exacta por CODIGO_UNICO ni auxiliar fuerte en el listado existente."

        base.update(
            {
                "ESTADO_CRUCE_INDICES": estado,
                "MOTIVO_ESTADO_CRUCE": motivo,
                "CODIGO_UNICO_MATCH": codigo_match,
                "TIPO_MATCH": tipo_match,
                "CONFIANZA_MATCH": confianza,
                "OBSERVACION_AUDITORIA": observacion,
                "REQUIERE_CREACION_MANUAL": requiere,
            }
        )
        rows.append(base)
    return pd.DataFrame(rows).fillna("")


def build_programas_faltantes(cruce_df: pd.DataFrame) -> pd.DataFrame:
    faltantes = cruce_df[cruce_df["ESTADO_CRUCE_INDICES"] == "FALTA_CREAR_EN_INDICES"].copy()
    faltantes = faltantes.sort_values(["SEDE_CRUCE_NORMALIZADA", "COD_CARRERA_DERIVADO", "VERSION_DERIVADA", "CODIGO_UNICO"])
    faltantes["Sede"] = faltantes["SEDE_FORMULARIO"]
    faltantes["Tipo de Carrera"] = faltantes["TIPO_DE_CARRERA_FORMULARIO"]
    faltantes["Nombre"] = faltantes["NOMBRE_CARRERA"]
    faltantes["Mención o Especialidad"] = ""
    faltantes["Año de inicio de Actividades"] = "REVISAR"
    faltantes["Campus"] = faltantes["CAMPUS_FORMULARIO"]
    faltantes["Horario"] = faltantes["HORARIO_FORMULARIO"]
    faltantes["Estado"] = "Programa o Carrera Nueva(o)"
    faltantes["Tipo Programa"] = faltantes["TIPO_PROGRAMA_FORMULARIO"]
    faltantes["Detalle del Tipo de Programa (especiales)"] = faltantes["DETALLE_TIPO_PROGRAMA_ESPECIAL_FORMULARIO"]
    faltantes["Modalidad del Programa"] = faltantes["MODALIDAD_DEL_PROGRAMA_FORMULARIO"]
    faltantes["Área del Conocimiento"] = "REVISAR"
    faltantes["Sub Área"] = "REVISAR"
    faltantes["Carrera Genérica"] = "REVISAR"
    faltantes["Título que otorga el programa"] = "REVISAR"
    faltantes["Grado Académico que otorga el programa"] = "REVISAR"
    faltantes["Dependencia"] = "REVISAR"
    faltantes["Régimen"] = "REVISAR"
    faltantes["Duración del Programa en Semestres"] = faltantes["DURACION_ESTUDIOS"]
    faltantes["Ingreso Directo"] = "REVISAR"
    faltantes["Ingreso desde un plan Común"] = "REVISAR"
    faltantes["Ingreso desde Bachillerato"] = "REVISAR"
    faltantes["Otro Tipo de Ingreso"] = "REVISAR"
    faltantes["Ingreso otro"] = "REVISAR"
    faltantes["Observaciones"] = faltantes.apply(
        lambda row: (
            "Crear segun matriz PROMEDIOSDEALUMNOS_7804.xlsx / hoja matriz; "
            "no encontrado en listado existente INDICES por CODIGO_UNICO; "
            f"sede={row['SEDE_FORMULARIO']}; tipo_carrera={row['TIPO_DE_CARRERA_FORMULARIO']}; "
            f"modalidad={row['MODALIDAD_DEL_PROGRAMA_FORMULARIO']}; horario={row['HORARIO_FORMULARIO']}"
        ),
        axis=1,
    )
    faltantes["NOMBRE_CARRERA_ORIGINAL"] = faltantes["NOMBRE_CARRERA"]
    faltantes["MODALIDAD_ORIGINAL"] = faltantes["MODALIDAD"]
    faltantes["TIPO_PLAN_CARRERA_ORIGINAL"] = faltantes["TIPO_PLAN_CARRERA"]
    faltantes["CARACT_PLAN_ESPECIAL_ORIGINAL"] = faltantes["CARACT_PLAN_ESPECIAL"]
    faltantes["REVISADO_MANUALMENTE"] = "NO"
    ordered = FORM_COLUMNS + TRACE_COLUMNS
    return faltantes[ordered].copy()


def build_expected_codes_audit(offer_df: pd.DataFrame, existing_df: pd.DataFrame, cross_df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    for code, expected_name in EXPECTED_CODES:
        offer_row = offer_df[offer_df["CODIGO_UNICO"] == code]
        existing_row = existing_df[existing_df["CODIGO_UNICO"] == code]
        cross_row = cross_df[cross_df["CODIGO_UNICO"] == code]
        rows.append(
            {
                "CODIGO_UNICO": code,
                "NOMBRE_ESPERADO": expected_name,
                "EN_OFERTA_MATRIZ_NORMALIZADA": "SI" if not offer_row.empty else "NO",
                "NOMBRE_EN_MATRIZ": clean_cell(offer_row["NOMBRE_CARRERA"].iloc[0]) if not offer_row.empty else "",
                "EN_PROGRAMAS_EXISTENTES_NORMALIZADOS": "SI" if not existing_row.empty else "NO",
                "ESTADO_CRUCE_INDICES": clean_cell(cross_row["ESTADO_CRUCE_INDICES"].iloc[0]) if not cross_row.empty else "",
                "OBSERVACION": clean_cell(cross_row["OBSERVACION_AUDITORIA"].iloc[0]) if not cross_row.empty else "Codigo esperado ausente del cruce.",
            }
        )
    return pd.DataFrame(rows)


def build_duplicates_audit(offer_df: pd.DataFrame, existing_df: pd.DataFrame, faltantes_df: pd.DataFrame) -> pd.DataFrame:
    existing_dup_codes = existing_df.loc[(existing_df["CODIGO_UNICO"] != "") & existing_df["CODIGO_UNICO"].duplicated(keep=False), "CODIGO_UNICO"].sort_values().unique().tolist()
    return pd.DataFrame(
        [
            {
                "TABLA": "OFERTA_MATRIZ_NORMALIZADA",
                "CAMPO": "CODIGO_UNICO",
                "TOTAL_FILAS": str(len(offer_df)),
                "TOTAL_DUPLICADOS": str(int(offer_df["CODIGO_UNICO"].duplicated().sum())),
                "EJEMPLOS": "",
                "OBSERVACION": "No se esperan duplicados por CODIGO_UNICO en matriz.",
            },
            {
                "TABLA": "PROGRAMAS_EXISTENTES_NORMALIZADOS",
                "CAMPO": "CODIGO_UNICO",
                "TOTAL_FILAS": str(len(existing_df)),
                "TOTAL_DUPLICADOS": str(len(existing_dup_codes)),
                "EJEMPLOS": " | ".join(existing_dup_codes),
                "OBSERVACION": "Duplicados en existentes no bloquean exact match, pero quedan auditados.",
            },
            {
                "TABLA": "PROGRAMAS_FALTANTES_PARA_CREAR",
                "CAMPO": "CODIGO_UNICO",
                "TOTAL_FILAS": str(len(faltantes_df)),
                "TOTAL_DUPLICADOS": str(int(faltantes_df["CODIGO_UNICO"].duplicated().sum() if not faltantes_df.empty else 0)),
                "EJEMPLOS": "",
                "OBSERVACION": "La hoja principal no debe contener duplicados por CODIGO_UNICO.",
            },
        ]
    )


def build_codigo_unico_audit(offer_df: pd.DataFrame, existing_df: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "TABLA": "OFERTA_MATRIZ_NORMALIZADA",
                "TOTAL_FILAS": str(len(offer_df)),
                "CODIGO_UNICO_VACIO": str(int((offer_df["CODIGO_UNICO"] == "").sum())),
                "FORMATO_OK": str(int((offer_df["FORMATO_CODIGO_UNICO"] == "OK").sum())),
                "FORMATO_REVISAR": str(int((offer_df["FORMATO_CODIGO_UNICO"] != "OK").sum())),
                "DUPLICADOS": str(int(offer_df["CODIGO_UNICO"].duplicated().sum())),
                "VALIDACION_SEDE_OK": str(int((offer_df["VALIDACION_SEDE_CODIGO_UNICO"] == "OK").sum())),
                "VALIDACION_JORNADA_OK": str(int((offer_df["VALIDACION_JORNADA_CODIGO_UNICO"] == "OK").sum())),
            },
            {
                "TABLA": "PROGRAMAS_EXISTENTES_NORMALIZADOS",
                "TOTAL_FILAS": str(len(existing_df)),
                "CODIGO_UNICO_VACIO": str(int((existing_df["CODIGO_UNICO"] == "").sum())),
                "FORMATO_OK": str(int((existing_df["FORMATO_CODIGO_UNICO"] == "OK").sum())),
                "FORMATO_REVISAR": str(int(((existing_df["CODIGO_UNICO"] != "") & (existing_df["FORMATO_CODIGO_UNICO"] != "OK")).sum())),
                "DUPLICADOS": str(int(existing_df.loc[existing_df["CODIGO_UNICO"] != "", "CODIGO_UNICO"].duplicated().sum())),
                "VALIDACION_SEDE_OK": "N/A",
                "VALIDACION_JORNADA_OK": "N/A",
            },
        ]
    )


def build_auditoria_columnas(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    requirements = {
        "OFERTA_MATRIZ_NORMALIZADA": MATRIX_COLUMNS + ["COD_CARRERA_DERIVADO", "NOMBRE_CARRERA_NORMALIZADO", "ALCANCE_INDICES_CREAR_PROGRAMA"],
        "PROGRAMAS_EXISTENTES_NORMALIZADOS": ["Codigo SIES", "CODIGO_UNICO", "NOMBRE_PROGRAMA_NORMALIZADO", "SEDE_CRUCE_NORMALIZADA"],
        "CRUCE_MATRIZ_VS_INDICES": ["ESTADO_CRUCE_INDICES", "MOTIVO_ESTADO_CRUCE", "TIPO_MATCH", "CONFIANZA_MATCH", "REQUIERE_CREACION_MANUAL"],
        "PROGRAMAS_FALTANTES_PARA_CREAR": FORM_COLUMNS + TRACE_COLUMNS,
    }
    for table_name, columns in requirements.items():
        current = tables[table_name]
        for column in columns:
            rows.append(
                {
                    "TABLA": table_name,
                    "COLUMNA": column,
                    "EXISTE": "SI" if column in current.columns else "NO",
                    "TOTAL_COLUMNAS_TABLA": str(len(current.columns)),
                    "OBSERVACION": "columna requerida por auditoria",
                }
            )
    return pd.DataFrame(rows)


def build_resumen(
    paths: Paths,
    sheet_meta: list[dict[str, Any]],
    offer_df: pd.DataFrame,
    existing_df: pd.DataFrame,
    cross_df: pd.DataFrame,
    faltantes_df: pd.DataFrame,
    dudosos_df: pd.DataFrame,
    ya_existen_df: pd.DataFrame,
    no_crear_df: pd.DataFrame,
    dictamen: str,
    warnings: list[str],
    errors: list[str],
    campos_con_revisar: list[str],
    git_snapshot: dict[str, str],
) -> pd.DataFrame:
    rows = [
        {"Campo": "DICTAMEN_FINAL", "Valor": dictamen},
        {"Campo": "FUENTE_PRINCIPAL", "Valor": "PROMEDIOSDEALUMNOS_7804.xlsx / hoja matriz"},
        {"Campo": "FUENTE_EXISTENTES_INDICES", "Valor": str(paths.existing_path)},
        {"Campo": "TOTAL_OFERTA_MATRIZ", "Valor": str(len(offer_df))},
        {"Campo": "TOTAL_VIGENTES", "Valor": str(int((offer_df["VIGENCIA"] == "1").sum()))},
        {"Campo": "TOTAL_YA_EXISTE_EN_INDICES", "Valor": str(len(ya_existen_df))},
        {"Campo": "TOTAL_FALTA_CREAR_EN_INDICES", "Valor": str(len(faltantes_df))},
        {"Campo": "TOTAL_DUDOSO_REVISAR_MANUAL", "Valor": str(len(dudosos_df))},
        {"Campo": "TOTAL_NO_CREAR", "Valor": str(len(no_crear_df))},
        {"Campo": "ARCHIVO_EXCEL_FINAL", "Valor": str(paths.excel_out)},
        {"Campo": "HOJA_PRINCIPAL_PARA_TRABAJO_MANUAL", "Valor": "PROGRAMAS_FALTANTES_PARA_CREAR"},
        {"Campo": "CAMPOS_FORMULARIO_COMPLETADOS", "Valor": str(len(FORM_COLUMNS))},
        {"Campo": "CAMPOS_CON_REVISAR", "Valor": ", ".join(campos_con_revisar) if campos_con_revisar else "ninguno"},
        {"Campo": "ERRORES_BLOQUEANTES", "Valor": " | ".join(errors) if errors else "ninguno"},
        {"Campo": "ADVERTENCIAS_NO_BLOQUEANTES", "Valor": " | ".join(warnings) if warnings else "ninguna"},
        {"Campo": "ESTADO_GIT_FINAL", "Valor": f"branch={git_snapshot['branch']} status_short={git_snapshot['status_short']} divergence={git_snapshot['divergence']}"},
        {"Campo": "RUTA_SCRIPT", "Valor": paths.script_path},
        {"Campo": "HOJAS_PROMEDIOS_DETECTADAS", "Valor": ", ".join(meta['sheet'] for meta in sheet_meta)},
    ]
    return pd.DataFrame(rows)


def build_bitacora(paths: Paths, git_snapshot: dict[str, str], validations: list[str]) -> pd.DataFrame:
    entries = [
        {"PASO": "01_BASELINE", "DETALLE": "Repo y entorno verificados externamente antes del script", "VALOR": git_snapshot["branch"]},
        {"PASO": "02_FUENTE_PRINCIPAL", "DETALLE": "Fuente de verdad de oferta", "VALOR": str(paths.matrix_path)},
        {"PASO": "03_FUENTE_EXISTENTES", "DETALLE": "Listado de programas existentes", "VALOR": str(paths.existing_path)},
        {"PASO": "04_NORMALIZACION", "DETALLE": "Texto a mayusculas sin tildes y parseo de CODIGO_UNICO", "VALOR": "OK"},
        {"PASO": "05_CRUCE", "DETALLE": RULES_CROSS, "VALOR": "OK"},
        {"PASO": "06_SALIDAS", "DETALLE": "Excel/CSV/Markdown/JSON timestamped", "VALOR": str(paths.excel_out)},
    ]
    entries.extend({"PASO": "VALIDACION", "DETALLE": item, "VALOR": "OK"} for item in validations)
    return pd.DataFrame(entries)


def attach_metadata(df: pd.DataFrame, generated_at: str, source_label: str, source_path: str, rules: str) -> pd.DataFrame:
    out = df.copy()
    out["META_GENERADO"] = generated_at
    out["META_TOTAL_FILAS_HOJA"] = str(len(df))
    out["META_FUENTE_DATOS"] = source_label
    out["META_RUTA_FUENTE"] = source_path
    out["META_REGLAS_APLICADAS"] = rules
    return out


def adjust_widths(ws: Any) -> None:
    for col_idx in range(1, ws.max_column + 1):
        values = [clean_cell(ws.cell(row=row_idx, column=col_idx).value) for row_idx in range(1, ws.max_row + 1)]
        width = min(max(max((len(value) for value in values), default=0) + 2, 12), 70)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_sheet(ws: Any, title: str) -> None:
    header_map = {}
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_map[clean_cell(cell.value)] = col_idx
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = clean_cell(cell.value)
            if value == "REVISAR":
                cell.fill = WARN_FILL
            elif value == "OK":
                cell.fill = OK_FILL
            elif value in {"NO", "ERROR_BLOQUEANTE"}:
                cell.fill = ERROR_FILL
            elif value in {"SI", "FALTA_CREAR_EN_INDICES", "DUDOSO_REVISAR_MANUAL"}:
                cell.fill = INFO_FILL
    estado_col = header_map.get("ESTADO_CRUCE_INDICES")
    if estado_col:
        for row_idx in range(2, ws.max_row + 1):
            value = clean_cell(ws.cell(row_idx, estado_col).value)
            if value == "YA_EXISTE_EN_INDICES":
                ws.cell(row_idx, estado_col).fill = OK_FILL
            elif value == "FALTA_CREAR_EN_INDICES":
                ws.cell(row_idx, estado_col).fill = INFO_FILL
            elif value == "DUDOSO_REVISAR_MANUAL":
                ws.cell(row_idx, estado_col).fill = WARN_FILL
            elif value.startswith("NO_CREAR"):
                ws.cell(row_idx, estado_col).fill = ERROR_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    adjust_widths(ws)


def add_dataframe_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    if df.empty:
        ws.append(list(df.columns))
    else:
        for row in dataframe_to_rows(df.fillna(""), index=False, header=True):
            ws.append(row)
    style_sheet(ws, title)


def collect_validations(
    paths: Paths,
    offer_df: pd.DataFrame,
    existing_df: pd.DataFrame,
    cross_df: pd.DataFrame,
    faltantes_df: pd.DataFrame,
    expected_df: pd.DataFrame,
) -> tuple[list[str], list[str]]:
    validations: list[str] = []
    errors: list[str] = []

    if paths.matrix_path.exists():
        validations.append("PROMEDIOSDEALUMNOS_7804.xlsx existe")
    else:
        errors.append("PROMEDIOSDEALUMNOS_7804.xlsx no existe")

    if len(offer_df) > 0:
        validations.append("hoja matriz existe y contiene filas")
    else:
        errors.append("hoja matriz sin filas")

    if int((offer_df["CODIGO_UNICO"] == "").sum()) == 0:
        validations.append("CODIGO_UNICO existe y no esta vacio en matriz")
    else:
        errors.append("existen CODIGO_UNICO vacios en matriz")

    if int((offer_df["FORMATO_CODIGO_UNICO"] == "OK").sum()) == len(offer_df):
        validations.append("todo CODIGO_UNICO en matriz cumple formato esperado")
    else:
        errors.append("existen CODIGO_UNICO con formato invalido en matriz")

    if int(offer_df["CODIGO_UNICO"].duplicated().sum()) == 0:
        validations.append("no hay duplicados inesperados por CODIGO_UNICO en matriz")
    else:
        errors.append("hay duplicados por CODIGO_UNICO en matriz")

    if int((offer_df["VALIDACION_SEDE_CODIGO_UNICO"] == "OK").sum()) == len(offer_df) and int((offer_df["VALIDACION_JORNADA_CODIGO_UNICO"] == "OK").sum()) == len(offer_df):
        validations.append("los derivados de CODIGO_UNICO son coherentes con sede y jornada")
    else:
        errors.append("hay incoherencias entre CODIGO_UNICO y sede/jornada")

    if len(existing_df) > 0:
        validations.append("existe base de programas ya existentes en INDICES/SIES")
    else:
        errors.append("no se identifico base de programas existentes")

    if len(cross_df) == len(offer_df):
        validations.append("el cruce por CODIGO_UNICO se ejecuto para toda la matriz")
    else:
        errors.append("el cruce no cubre toda la matriz")

    if set(faltantes_df["ESTADO_CRUCE_INDICES"].unique()) <= {"FALTA_CREAR_EN_INDICES"}:
        validations.append("PROGRAMAS_FALTANTES_PARA_CREAR contiene solo FALTA_CREAR_EN_INDICES")
    else:
        errors.append("PROGRAMAS_FALTANTES_PARA_CREAR mezcla estados no permitidos")

    if int((faltantes_df["CODIGO_UNICO"] == "").sum()) == 0:
        validations.append("no hay CODIGO_UNICO vacio en PROGRAMAS_FALTANTES_PARA_CREAR")
    else:
        errors.append("hay CODIGO_UNICO vacio en PROGRAMAS_FALTANTES_PARA_CREAR")

    if int(faltantes_df["CODIGO_UNICO"].duplicated().sum()) == 0:
        validations.append("no hay duplicados en PROGRAMAS_FALTANTES_PARA_CREAR")
    else:
        errors.append("hay duplicados en PROGRAMAS_FALTANTES_PARA_CREAR")

    missing_form_columns = [column for column in FORM_COLUMNS if column not in faltantes_df.columns]
    if not missing_form_columns:
        validations.append("todos los campos de pantalla estan presentes")
    else:
        errors.append(f"faltan columnas de formulario: {missing_form_columns}")

    missing_expected = expected_df[expected_df["EN_OFERTA_MATRIZ_NORMALIZADA"] != "SI"]
    if missing_expected.empty:
        validations.append("todos los codigos esperados del usuario aparecen en OFERTA_MATRIZ_NORMALIZADA")
    else:
        errors.append("faltan codigos esperados del usuario en OFERTA_MATRIZ_NORMALIZADA")

    return validations, errors


def compute_campos_con_revisar(faltantes_df: pd.DataFrame) -> list[str]:
    if faltantes_df.empty:
        return []
    columns: list[str] = []
    for column in FORM_COLUMNS:
        if column in faltantes_df.columns and faltantes_df[column].astype(str).str.contains("REVISAR", regex=False).any():
            columns.append(column)
    return columns


def build_markdown(
    paths: Paths,
    dictamen: str,
    resumen_df: pd.DataFrame,
    expected_df: pd.DataFrame,
    faltantes_df: pd.DataFrame,
    dudosos_df: pd.DataFrame,
    no_crear_df: pd.DataFrame,
    validations: list[str],
    warnings: list[str],
    errors: list[str],
    campos_con_revisar: list[str],
) -> str:
    resumen = {row["Campo"]: row["Valor"] for _, row in resumen_df.iterrows()}
    expected_rows = expected_df[["CODIGO_UNICO", "NOMBRE_ESPERADO", "EN_PROGRAMAS_EXISTENTES_NORMALIZADOS", "ESTADO_CRUCE_INDICES"]].to_dict(orient="records")
    expected_md = "\n".join(
        f"- {row['CODIGO_UNICO']} | {row['NOMBRE_ESPERADO']} | existentes={row['EN_PROGRAMAS_EXISTENTES_NORMALIZADOS']} | estado={row['ESTADO_CRUCE_INDICES']}"
        for row in expected_rows
    )
    validations_md = "\n".join(f"- {item}" for item in validations)
    warnings_md = "\n".join(f"- {item}" for item in warnings) if warnings else "- ninguna"
    errors_md = "\n".join(f"- {item}" for item in errors) if errors else "- ninguno"
    return f"""# CNED Programas Faltantes para Crear en INDICES

## Dictamen

{dictamen}

## Fuentes usadas

- FUENTE_PRINCIPAL: PROMEDIOSDEALUMNOS_7804.xlsx / hoja matriz
- FUENTE_EXISTENTES_INDICES: {paths.existing_path}
- Script ejecutado: {paths.script_path}

## Conteos

- TOTAL_OFERTA_MATRIZ: {resumen['TOTAL_OFERTA_MATRIZ']}
- TOTAL_VIGENTES: {resumen['TOTAL_VIGENTES']}
- TOTAL_YA_EXISTE_EN_INDICES: {resumen['TOTAL_YA_EXISTE_EN_INDICES']}
- TOTAL_FALTA_CREAR_EN_INDICES: {resumen['TOTAL_FALTA_CREAR_EN_INDICES']}
- TOTAL_DUDOSO_REVISAR_MANUAL: {resumen['TOTAL_DUDOSO_REVISAR_MANUAL']}
- TOTAL_NO_CREAR: {resumen['TOTAL_NO_CREAR']}

## Reglas de cruce

- {RULES_CROSS}
- Toda coincidencia exacta por CODIGO_UNICO se clasifica YA_EXISTE_EN_INDICES.
- Coincidencias auxiliares fuertes quedan como DUDOSO_REVISAR_MANUAL.
- Sin coincidencia exacta ni auxiliar fuerte, el estado es FALTA_CREAR_EN_INDICES.
- Todo campo de formulario sin evidencia directa queda como REVISAR.

## Programas existentes, faltantes y dudosos

- YA_EXISTEN_EN_INDICES: {len(expected_df[expected_df['ESTADO_CRUCE_INDICES'] == 'YA_EXISTE_EN_INDICES'])} codigos esperados con match exacto en la auditoria de codigos esperados.
- PROGRAMAS_FALTANTES_PARA_CREAR: {len(faltantes_df)} filas.
- DUDOSOS_REVISAR_MANUAL: {len(dudosos_df)} filas.
- NO_CREAR: {len(no_crear_df)} filas.

## Validaciones

{validations_md}

## Errores bloqueantes

{errors_md}

## Advertencias no bloqueantes

{warnings_md}

## Auditoria de codigos esperados

{expected_md}

## Archivo final de trabajo

- ARCHIVO_EXCEL_FINAL: {paths.excel_out}
- CSV principal: {paths.csv_out}
- JSON resumen: {paths.json_out}
- HOJA_PRINCIPAL_PARA_TRABAJO_MANUAL: PROGRAMAS_FALTANTES_PARA_CREAR

## Campos del formulario

- CAMPOS_FORMULARIO_COMPLETADOS: {len(FORM_COLUMNS)}
- CAMPOS_CON_REVISAR: {', '.join(campos_con_revisar) if campos_con_revisar else 'ninguno'}

## Proximos pasos

- usar hoja PROGRAMAS_FALTANTES_PARA_CREAR para completar pantalla Crear Programa
- revisar hoja DUDOSOS_REVISAR_MANUAL antes de crear cualquier programa ambiguo
- no usar la base historica reducida de 59 filas como universo principal para este flujo
"""


def post_generation_audit(paths: Paths) -> tuple[list[str], list[str], dict[str, Any]]:
    validations: list[str] = []
    errors: list[str] = []
    wb = load_workbook(paths.excel_out, read_only=True, data_only=False)
    if all(sheet in wb.sheetnames for sheet in REQUIRED_SHEETS):
        validations.append("el Excel final existe y contiene todas las hojas obligatorias")
    else:
        missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in wb.sheetnames]
        errors.append(f"faltan hojas obligatorias en Excel final: {missing}")

    faltantes_df = pd.read_excel(paths.excel_out, sheet_name="PROGRAMAS_FALTANTES_PARA_CREAR", dtype=str).fillna("")
    csv_df = pd.read_csv(paths.csv_out, sep=";", dtype=str, keep_default_na=False)
    if len(csv_df) == len(faltantes_df):
        validations.append("CSV principal tiene las mismas filas que PROGRAMAS_FALTANTES_PARA_CREAR")
    else:
        errors.append("CSV principal no coincide en numero de filas con PROGRAMAS_FALTANTES_PARA_CREAR")

    required_form = [column for column in FORM_COLUMNS if column not in faltantes_df.columns]
    if not required_form:
        validations.append("PROGRAMAS_FALTANTES_PARA_CREAR contiene todas las columnas del formulario")
    else:
        errors.append(f"faltan columnas del formulario en Excel final: {required_form}")

    if int((faltantes_df.get("CODIGO_UNICO", pd.Series(dtype=str)) == "").sum()) == 0:
        validations.append("no hay CODIGO_UNICO vacio en la hoja principal")
    else:
        errors.append("hay CODIGO_UNICO vacio en la hoja principal")

    if int(faltantes_df.get("CODIGO_UNICO", pd.Series(dtype=str)).duplicated().sum()) == 0:
        validations.append("no hay duplicados por CODIGO_UNICO en la hoja principal")
    else:
        errors.append("hay duplicados por CODIGO_UNICO en la hoja principal")

    if set(faltantes_df.get("ESTADO_CRUCE_INDICES", pd.Series(dtype=str)).unique()) <= {"FALTA_CREAR_EN_INDICES"}:
        validations.append("la hoja principal solo contiene FALTA_CREAR_EN_INDICES")
    else:
        errors.append("la hoja principal contiene estados distintos a FALTA_CREAR_EN_INDICES")

    snapshot = {
        "excel_exists": paths.excel_out.exists(),
        "csv_exists": paths.csv_out.exists(),
        "md_exists": paths.md_out.exists(),
        "json_exists": paths.json_out.exists(),
        "sheet_names": wb.sheetnames,
        "faltantes_rows": len(faltantes_df),
        "csv_rows": len(csv_df),
    }
    return validations, errors, snapshot


def determine_dictamen(errors: list[str], faltantes_df: pd.DataFrame, dudosos_df: pd.DataFrame) -> str:
    if errors:
        return "ERROR_BLOQUEANTE"
    if len(faltantes_df) > 0:
        return "OK_PROGRAMAS_FALTANTES_IDENTIFICADOS_PARA_CREAR"
    if len(dudosos_df) > 0:
        return "REQUIERE_REVISION_MANUAL"
    return "SIN_FALTANTES_TODO_EXISTE"


def main() -> int:
    paths = build_paths()
    paths.resultados.mkdir(parents=True, exist_ok=True)

    warnings: list[str] = []
    errors: list[str] = []

    manual_tsv_stat_before = paths.manual_tsv_path.stat() if paths.manual_tsv_path.exists() else None
    git_snapshot = {
        "branch": git_output(paths.repo, "branch", "--show-current"),
        "status_short": git_output(paths.repo, "status", "--short"),
        "head": git_output(paths.repo, "rev-parse", "--short", "HEAD"),
        "remote_head": git_output(paths.repo, "rev-parse", "--short", "origin/clean/pes-ready-final"),
        "divergence": git_output(paths.repo, "log", "--oneline", "--left-right", "HEAD...origin/clean/pes-ready-final"),
    }

    try:
        matrix_raw, sheet_meta, workbook_sheets = read_matrix(paths)
        existing_raw = read_existing(paths)
    except Exception as exc:
        error_md = paths.resultados / "ERROR_NO_SE_IDENTIFICA_LISTADO_EXISTENTE_INDICES.md"
        error_md.write_text(f"# ERROR_NO_SE_IDENTIFICA_LISTADO_EXISTENTE_INDICES\n\n{exc}\n", encoding="utf-8")
        print(f"ERROR_BLOQUEANTE: {exc}")
        print(f"Informe generado: {error_md}")
        return 1

    offer_normalized = build_offer_matrix_normalized(matrix_raw, paths)
    existing_normalized = build_existing_normalized(existing_raw, paths)
    dictionaries, dictionary_df, dictionary_warnings = build_dictionaries(offer_normalized, existing_normalized)
    warnings.extend(dictionary_warnings)

    offer_classified = build_offer_matrix_classified(offer_normalized, dictionaries)
    cross_df = build_cross(offer_classified, existing_normalized)
    dudosos_df = cross_df[cross_df["ESTADO_CRUCE_INDICES"] == "DUDOSO_REVISAR_MANUAL"].copy()
    ya_existen_df = cross_df[cross_df["ESTADO_CRUCE_INDICES"] == "YA_EXISTE_EN_INDICES"].copy()
    no_crear_df = cross_df[cross_df["ESTADO_CRUCE_INDICES"].isin(["NO_CREAR_POR_NO_VIGENTE", "NO_CREAR_POR_FUERA_DE_ALCANCE"])].copy()
    faltantes_df = build_programas_faltantes(cross_df)
    expected_df = build_expected_codes_audit(offer_classified, existing_normalized, cross_df)

    validations, pre_errors = collect_validations(paths, offer_classified, existing_normalized, cross_df, faltantes_df, expected_df)
    errors.extend(pre_errors)

    if expected_df[expected_df["EN_OFERTA_MATRIZ_NORMALIZADA"] != "SI"].shape[0] > 0:
        errors.append("La auditoria de codigos esperados indica que la fuente no corresponde a la matriz correcta.")

    campos_con_revisar = compute_campos_con_revisar(faltantes_df)
    if campos_con_revisar:
        warnings.append(f"Campos con REVISAR en hoja principal: {', '.join(campos_con_revisar)}")
    if len(dudosos_df) > 0:
        warnings.append(f"Existen {len(dudosos_df)} registros dudosos que requieren revision manual antes de crear programas ambiguos.")
    if git_snapshot["divergence"]:
        warnings.append("La rama local sigue divergida respecto de origin/clean/pes-ready-final; los artefactos generados son locales hasta nueva decision Git.")

    dictamen = determine_dictamen(errors, faltantes_df, dudosos_df)

    audit_duplicates_df = build_duplicates_audit(offer_classified, existing_normalized, faltantes_df)
    audit_codigo_df = build_codigo_unico_audit(offer_classified, existing_normalized)
    tables_for_audit = {
        "OFERTA_MATRIZ_NORMALIZADA": offer_classified,
        "PROGRAMAS_EXISTENTES_NORMALIZADOS": existing_normalized,
        "CRUCE_MATRIZ_VS_INDICES": cross_df,
        "PROGRAMAS_FALTANTES_PARA_CREAR": faltantes_df,
    }
    audit_columnas_df = build_auditoria_columnas(tables_for_audit)
    resumen_df = build_resumen(
        paths,
        sheet_meta,
        offer_classified,
        existing_normalized,
        cross_df,
        faltantes_df,
        dudosos_df,
        ya_existen_df,
        no_crear_df,
        dictamen,
        warnings,
        errors,
        campos_con_revisar,
        git_snapshot,
    )
    bitacora_df = build_bitacora(paths, git_snapshot, validations)

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    workbook_tables = {
        "RESUMEN_EJECUTIVO": attach_metadata(resumen_df, generated_at, "resumen del proceso", str(paths.matrix_path), RULES_CROSS),
        "OFERTA_MATRIZ_NORMALIZADA": attach_metadata(offer_classified, generated_at, "PROMEDIOSDEALUMNOS_7804.xlsx / matriz", str(paths.matrix_path), "normalizacion de matriz + derivados de CODIGO_UNICO"),
        "PROGRAMAS_EXISTENTES_NORMALIZADOS": attach_metadata(existing_normalized, generated_at, "listado_referencia_cned.tsv", str(paths.existing_path), "normalizacion de existentes + derivados de CODIGO_UNICO"),
        "CRUCE_MATRIZ_VS_INDICES": attach_metadata(cross_df, generated_at, "matriz vs existentes", f"{paths.matrix_path} | {paths.existing_path}", RULES_CROSS),
        "PROGRAMAS_FALTANTES_PARA_CREAR": attach_metadata(faltantes_df, generated_at, "filtrado de CRUCE_MATRIZ_VS_INDICES", str(paths.matrix_path), "solo FALTA_CREAR_EN_INDICES; campos sin evidencia=REVISAR"),
        "DUDOSOS_REVISAR_MANUAL": attach_metadata(dudosos_df, generated_at, "filtrado de CRUCE_MATRIZ_VS_INDICES", str(paths.matrix_path), "coincidencias auxiliares fuertes sin match exacto"),
        "YA_EXISTEN_EN_INDICES": attach_metadata(ya_existen_df, generated_at, "filtrado de CRUCE_MATRIZ_VS_INDICES", str(paths.existing_path), "coincidencia exacta por CODIGO_UNICO"),
        "NO_CREAR": attach_metadata(no_crear_df, generated_at, "filtrado de CRUCE_MATRIZ_VS_INDICES", str(paths.matrix_path), "no vigente o fuera de alcance"),
        "DICCIONARIO_CAMPOS": attach_metadata(dictionary_df, generated_at, "diccionarios de mapeo", f"{paths.matrix_path} | {paths.existing_path}", "mapeos de columnas y codigos"),
        "AUDITORIA_COLUMNAS": attach_metadata(audit_columnas_df, generated_at, "auditoria de columnas", str(paths.excel_out), "verificacion de columnas requeridas"),
        "AUDITORIA_DUPLICADOS": attach_metadata(audit_duplicates_df, generated_at, "auditoria de duplicados", f"{paths.matrix_path} | {paths.existing_path}", "duplicados por CODIGO_UNICO"),
        "AUDITORIA_CODIGO_UNICO": attach_metadata(audit_codigo_df, generated_at, "auditoria de codigo unico", f"{paths.matrix_path} | {paths.existing_path}", "formato, vacios y coherencia"),
        "AUDITORIA_CODIGOS_ESPERADOS": attach_metadata(expected_df, generated_at, "auditoria de codigos esperados", str(paths.matrix_path), "verificacion del universo esperado del usuario"),
        "BITACORA_PROCESO": attach_metadata(bitacora_df, generated_at, "bitacora del proceso", paths.script_path, "registro resumido de pasos y validaciones"),
    }

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in REQUIRED_SHEETS:
        add_dataframe_sheet(wb, sheet, workbook_tables[sheet])
    wb.save(paths.excel_out)

    faltantes_df.to_csv(paths.csv_out, sep=";", index=False, encoding="utf-8")
    markdown = build_markdown(paths, dictamen, resumen_df, expected_df, faltantes_df, dudosos_df, no_crear_df, validations, warnings, errors, campos_con_revisar)
    paths.md_out.write_text(markdown, encoding="utf-8")

    post_validations, post_errors, post_snapshot = post_generation_audit(paths)
    validations.extend(post_validations)
    errors.extend(post_errors)
    final_dictamen = determine_dictamen(errors, faltantes_df, dudosos_df)
    if final_dictamen != dictamen:
        resumen_df.loc[resumen_df["Campo"] == "DICTAMEN_FINAL", "Valor"] = final_dictamen
        markdown = build_markdown(paths, final_dictamen, resumen_df, expected_df, faltantes_df, dudosos_df, no_crear_df, validations, warnings, errors, campos_con_revisar)
        paths.md_out.write_text(markdown, encoding="utf-8")
        dictamen = final_dictamen

    manual_tsv_stat_after = paths.manual_tsv_path.stat() if paths.manual_tsv_path.exists() else None
    if manual_tsv_stat_before and manual_tsv_stat_after and manual_tsv_stat_before.st_mtime_ns == manual_tsv_stat_after.st_mtime_ns:
        validations.append("no se modifico el TSV manual")
    else:
        warnings.append("No fue posible probar con certeza la invariancia del TSV manual via mtime.")

    summary_json = {
        "timestamp": paths.timestamp,
        "dictamen_final": dictamen,
        "fuente_principal": "PROMEDIOSDEALUMNOS_7804.xlsx / hoja matriz",
        "fuente_existentes_indices": str(paths.existing_path),
        "paths": {
            "excel": str(paths.excel_out),
            "csv": str(paths.csv_out),
            "markdown": str(paths.md_out),
            "json": str(paths.json_out),
        },
        "counts": {
            "total_oferta_matriz": len(offer_classified),
            "total_vigentes": int((offer_classified["VIGENCIA"] == "1").sum()),
            "total_ya_existe_en_indices": len(ya_existen_df),
            "total_falta_crear_en_indices": len(faltantes_df),
            "total_dudoso_revisar_manual": len(dudosos_df),
            "total_no_crear": len(no_crear_df),
        },
        "validations": validations,
        "errores_bloqueantes": errors,
        "advertencias_no_bloqueantes": warnings,
        "campos_con_revisar": campos_con_revisar,
        "git": git_snapshot,
        "post_generation_audit": post_snapshot,
    }
    paths.json_out.write_text(json.dumps(summary_json, ensure_ascii=False, indent=2), encoding="utf-8")

    print("=" * 120)
    print("CNED - PROGRAMAS FALTANTES PARA CREAR EN INDICES")
    print("=" * 120)
    print(f"DICTAMEN_FINAL: {dictamen}")
    print("FUENTE_PRINCIPAL: PROMEDIOSDEALUMNOS_7804.xlsx / hoja matriz")
    print(f"FUENTE_EXISTENTES_INDICES: {paths.existing_path}")
    print(f"TOTAL_OFERTA_MATRIZ: {len(offer_classified)}")
    print(f"TOTAL_VIGENTES: {int((offer_classified['VIGENCIA'] == '1').sum())}")
    print(f"TOTAL_YA_EXISTE_EN_INDICES: {len(ya_existen_df)}")
    print(f"TOTAL_FALTA_CREAR_EN_INDICES: {len(faltantes_df)}")
    print(f"TOTAL_DUDOSO_REVISAR_MANUAL: {len(dudosos_df)}")
    print(f"TOTAL_NO_CREAR: {len(no_crear_df)}")
    print(f"ARCHIVO_EXCEL_FINAL: {paths.excel_out}")
    print("HOJA_PRINCIPAL_PARA_TRABAJO_MANUAL: PROGRAMAS_FALTANTES_PARA_CREAR")
    print(f"CAMPOS_FORMULARIO_COMPLETADOS: {len(FORM_COLUMNS)}")
    print(f"CAMPOS_CON_REVISAR: {', '.join(campos_con_revisar) if campos_con_revisar else 'ninguno'}")
    print(f"ERRORES_BLOQUEANTES: {' | '.join(errors) if errors else 'ninguno'}")
    print(f"ADVERTENCIAS_NO_BLOQUEANTES: {' | '.join(warnings) if warnings else 'ninguna'}")
    print(f"ESTADO_GIT_FINAL: branch={git_snapshot['branch']} | status_short={git_snapshot['status_short']} | divergence={git_snapshot['divergence']}")
    print("RECOMENDACION: abrir Excel final, usar hoja PROGRAMAS_FALTANTES_PARA_CREAR, crear manualmente en INDICES/CNED/SIES y revisar DUDOSOS_REVISAR_MANUAL antes de crear programas ambiguos")
    print("=" * 120)
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())