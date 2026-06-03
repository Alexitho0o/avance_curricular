#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Materializa el artefacto operativo CNED limpio en Excel, CSV y Markdown."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from io import StringIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


DICTAMEN = "BASE_CNED_LIMPIA_LISTA_PARA_USO_OPERATIVO"
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
HEADER_FONT = Font(color="FFFFFF", bold=True)
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="DDEBF7")
CODE_PATTERN = re.compile(r"^I162S(?P<sed>\d+)C(?P<car>\d+)J(?P<jor>\d+)V(?P<version>\d+)$")

BASE_CNED_LIMPIA_TSV = """Año	Cód. Institución	Nombre Institución	Nombre de la Sede	Comuna donde se imparte la carrera o programa	Cód. Carrera	Carrera Genérica	Nombre Programa	Horario	Tipo Programa	Tipo Carrera	IngresoDirecto	Año Inicio Actividades	Duración (en semestres)	Cód. Campus	Cód. Sede	Título	Código SIES	Pregrado/Posgrado	COD_SED_DERIVADO	COD_CAR_DERIVADO	JOR_DERIVADA	VERSION_DERIVADA	CLAVE_OPERATIVA	VALIDACION_CODIGO_UNICO	VALIDACION_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32513	Técnico en Automatización, Control automático y similares	TNS en Automatización y Control Industrial	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2015	5	2024001001	2024001	Técnico Nivel Sup. en Automatización y Control Industrial	I162S2C11J2V1	Pregrado	2	11	2	1	I162S2C11J2V1|32513|TNS EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6677	Ingeniería en Computación e Informática y similares	Ingeniería en Informática	Diurno	Programa Regular	Profesional	Ingreso Directo	1990	8		2024001	Ingeniero en Informática	I162S2C1J1V2	Pregrado	2	1	1	2	I162S2C1J1V2|6677|INGENIERÍA EN INFORMÁTICA|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6678	Ingeniería en Computación e Informática y similares	Ingeniería en Informática	Vespertino	Programa Regular	Profesional	Ingreso Directo	1990	8		2024001	Ingeniero en Informática	I162S2C1J2V4	Pregrado	2	1	2	4	I162S2C1J2V4|6678|INGENIERÍA EN INFORMÁTICA|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6680	Ingeniería en Conectividad y Redes	Ingeniería en Conectividad y Redes	Vespertino	Programa Regular	Profesional	Ingreso Directo	2002	8		2024001	Ingeniero en Conectividad y Redes	I162S2C3J2V4	Pregrado	2	3	2	4	I162S2C3J2V4|6680|INGENIERÍA EN CONECTIVIDAD Y REDES|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	38816	Ingeniería en Conectividad y Redes	Ingeniería en Conectividad y Redes	Otro	Programa Regular	Profesional	Ingreso Directo	2018	8	2024001001	2024001	Ingeniero en Conectividad y Redes	I162S2C3J4V2	Pregrado	2	3	4	2	I162S2C3J4V2|38816|INGENIERÍA EN CONECTIVIDAD Y REDES|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40713	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad	Diurno	Programa Regular	Profesional	Ingreso Directo	2019	8	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J1V1	Pregrado	2	46	1	1	I162S2C46J1V1|40713|INGENIERÍA EN CIBERSEGURIDAD|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40714	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad	Vespertino	Programa Regular	Profesional	Ingreso Directo	2019	8	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J2V1	Pregrado	2	46	2	1	I162S2C46J2V1|40714|INGENIERÍA EN CIBERSEGURIDAD|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	38817	Ingeniería en Computación e Informática y similares	Ingeniería en Informática	Otro	Programa Regular	Profesional	Ingreso Directo	2018	8	2024001001	2024001	Ingeniero en Informática	I162S2C46J4V2	Pregrado	2	46	4	2	I162S2C46J4V2|38817|INGENIERÍA EN INFORMÁTICA|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43928	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad	Otro	Programa Regular	Profesional	Ingreso Directo	2020	4	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J4V3	Pregrado	2	46	4	3	I162S2C46J4V3|43928|INGENIERÍA EN CIBERSEGURIDAD|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40717	Técnico en Análisis de Sistemas Informáticos	TNS en Ciberseguridad	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2019	5	2024001001	2024001	Técnico Nivel Sup. en Ciberseguridad	I162S2C47J1V1	Pregrado	2	47	1	1	I162S2C47J1V1|40717|TNS EN CIBERSEGURIDAD|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40711	Técnico en Análisis de Sistemas Informáticos	TNS en Ciberseguridad	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2019	5	2024001001	2024001	Técnico Nivel Sup. en Ciberseguridad	I162S2C47J2V1	Pregrado	2	47	2	1	I162S2C47J2V1|40711|TNS EN CIBERSEGURIDAD|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43931	Técnico en Análisis de Sistemas Informáticos	TNS en Ciberseguridad	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2020	5	2024001001	2024001	Técnico Nivel Sup. en Ciberseguridad	I162S2C47J4V1	Pregrado	2	47	4	1	I162S2C47J4V1|43931|TNS EN CIBERSEGURIDAD|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43935	Técnico Programador Computacional	TNS en Programación y Análisis de Sistemas	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2020	5	2024001001	2024001	Técnico Nivel Sup. en Programación y Análisis de Sistemas	I162S2C57J1V1	Pregrado	2	57	1	1	I162S2C57J1V1|43935|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43930	Técnico Programador Computacional	TNS en Programación y Análisis de Sistemas	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2020	5	2024001001	2024001	Técnico Nivel Sup. en Programación y Análisis de Sistemas	I162S2C57J2V1	Pregrado	2	57	2	1	I162S2C57J2V1|43930|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43932	Técnico Programador Computacional	TNS en Programación y Análisis de Sistemas	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2020	5	2024001001	2024001	Técnico Nivel Sup. en Programación Computacional y Análisis de Sistemas	I162S2C57J4V1	Pregrado	2	57	4	1	I162S2C57J4V1|43932|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6684	Técnico en Administración de Redes Computacionales	TNS en Conectividad y Redes	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2002	4		2024001	Técnico Nivel Sup. en Conectividad y Redes	I162S2C6J2V2	Pregrado	2	6	2	2	I162S2C6J2V2|6684|TNS EN CONECTIVIDAD Y REDES|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	36744	Técnico en Administración de Redes Computacionales	TNS en Conectividad y Redes	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2017	5	2024001001	2024001	Técnico Nivel Sup. en Conectividad y Redes	I162S2C6J4V2	Pregrado	2	6	4	2	I162S2C6J4V2|36744|TNS EN CONECTIVIDAD Y REDES|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51504	Ingeniería en Administración, Administración de Empresas y similares	Ingeniería en Administración de Empresas	Vespertino	Programa Regular	Profesional	Ingreso Directo	2024	8		2024001	Ingeniero en Administración de Empresas	I162S2C76J2V1	Pregrado	2	76	2	1	I162S2C76J2V1|51504|INGENIERÍA EN ADMINISTRACIÓN DE EMPRESAS|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51505	Ingeniería en Administración, Administración de Empresas y similares	Ingeniería en Administración de Empresas	Otro	Programa Regular	Profesional	Ingreso Directo	2024	8		2024001	Ingeniero en Administración de Empresas	I162S2C76J4V1	Pregrado	2	76	4	1	I162S2C76J4V1|51505|INGENIERÍA EN ADMINISTRACIÓN DE EMPRESAS|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51508	Ingeniería en Logística y similares	Ingeniería en Logística	Vespertino	Programa Regular	Profesional	Ingreso Directo	2024	8		2024001	Ingeniero en Logística	I162S2C77J2V1	Pregrado	2	77	2	1	I162S2C77J2V1|51508|INGENIERÍA EN LOGÍSTICA|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51509	Ingeniería en Logística y similares	Ingeniería en Logística	Otro	Programa Regular	Profesional	Ingreso Directo	2024	8		2024001	Ingeniero en Logística	I162S2C77J4V1	Pregrado	2	77	4	1	I162S2C77J4V1|51509|INGENIERÍA EN LOGÍSTICA|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51506	Técnico en Administración de Empresas	TNS en Administración de Empresas	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2024	5		2024001	Técnico Nivel Sup. en Administración de Empresas	I162S2C78J2V1	Pregrado	2	78	2	1	I162S2C78J2V1|51506|TNS EN ADMINISTRACIÓN DE EMPRESAS|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51510	Técnico en Administración de Empresas	TNS en Administración de Empresas	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2024	5		2024001	Técnico Nivel Sup. en Administración de Empresas	I162S2C78J4V1	Pregrado	2	78	4	1	I162S2C78J4V1|51510|TNS EN ADMINISTRACIÓN DE EMPRESAS|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51511	Técnico en Logística y similares	TNS en Logística	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2024	5		2024001	Técnico Nivel Sup. en Logística	I162S2C79J1V1	Pregrado	2	79	1	1	I162S2C79J1V1|51511|TNS EN LOGÍSTICA|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51512	Técnico en Logística y similares	TNS en Logística	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2024	5		2024001	Técnico Nivel Sup. en Logística	I162S2C79J2V1	Pregrado	2	79	2	1	I162S2C79J2V1|51512|TNS EN LOGÍSTICA|VESPERTINO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53717	Administración pública y similares	Administración Pública	Otro	Programa Regular	Profesional	Ingreso Directo	2025	8	2024001001	2024001	Administrador Público	I162S2C83J4V1	Pregrado	2	83	4	1	I162S2C83J4V1|53717|ADMINISTRACIÓN PÚBLICA|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53720	Técnico en Administración Pública	TNS en Administración Pública	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2025	5	2024001001	2024001	Técnico Nivel Sup. en Administración Pública	I162S2C84J4V1	Pregrado	2	84	4	1	I162S2C84J4V1|53720|TNS EN ADMINISTRACIÓN PÚBLICA|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53718	Técnico en Contabilidad, Auditoría y similares	TNS en Contabilidad General	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2025	5	2024001001	2024001	Técnico Nivel Sup. en Contabilidad General	I162S2C85J4V1	Pregrado	2	85	4	1	I162S2C85J4V1|53718|TNS EN CONTABILIDAD GENERAL|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53721	Contabilidad, Auditoría y similares	Auditoria	Otro	Programa Regular	Profesional	Ingreso Directo	2025	8	2024001001	2024001	Contador Auditor	I162S2C86J4V1	Pregrado	2	86	4	1	I162S2C86J4V1|53721|AUDITORIA|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53723	Ingeniería Ejecución Industrial y similares	Ingeniería Industrial (P.E.)	Otro	Programa Especial	Profesional	No es Ingreso Directo	2025	4	2024001001	2024001	Ingeniero Industrial	I162S2C87J4V1	Pregrado	2	87	4	1	I162S2C87J4V1|53723|INGENIERÍA INDUSTRIAL (P.E.)|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53722	Ingeniería en Computación e Informática y similares	Ingeniería en Ciencia de Datos	Otro	Programa Regular	Profesional	Ingreso Directo	2025	8	2024001001	2024001	Ingeniero en Ciencia de Datos	I162S2C88J4V1	Pregrado	2	88	4	1	I162S2C88J4V1|53722|INGENIERÍA EN CIENCIA DE DATOS|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53719	Técnico en Computación e Informática	TNS en Ciencia de Datos	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2025	5	2024001001	2024001	Técnico Nivel Sup. en Ciencia de Datos	I162S2C89J4V1	Pregrado	2	89	4	1	I162S2C89J4V1|53719|TNS EN CIENCIA DE DATOS|OTRO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53725	Técnico en Enfermería y Similares	TNS en Enfermería	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2025	5	2024001001	2024001	Técnico Nivel Sup. en Enfermería	I162S2C91J1V1	Pregrado	2	91	1	1	I162S2C91J1V1|53725|TNS EN ENFERMERÍA|DIURNO	OK	REVISAR_COD_CARRERA
2025	2024	IP SAN SEBASTIÁN	Santiago	Santiago	53726	Técnico en Enfermería y Similares	TNS en Enfermería	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2025	5	2024001001	2024001	Técnico Nivel Sup. en Enfermería	I162S2C91J2V1	Pregrado	2	91	2	1	I162S2C91J2V1|53726|TNS EN ENFERMERÍA|VESPERTINO	OK	REVISAR_COD_CARRERA
2024	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6679	Ingeniería en Conectividad y Redes	Ingeniería en Conectividad y Redes	Diurno	Programa Regular	Profesional	Ingreso Directo	2002	8		2024001	Ingeniería en Conectividad y Redes	I162S2C3J1V2	Pregrado	2	3	1	2	I162S2C3J1V2|6679|INGENIERÍA EN CONECTIVIDAD Y REDES|DIURNO	OK	REVISAR_COD_CARRERA
2024	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6683	Técnico en Administración de Redes Computacionales	TNS en Conectividad y Redes	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2002	4		2024001	Técnico Nivel Sup. en Conectividad y Redes	I162S2C6J1V2	Pregrado	2	6	1	2	I162S2C6J1V2|6683|TNS EN CONECTIVIDAD Y REDES|DIURNO	OK	REVISAR_COD_CARRERA
2024	2024	IP SAN SEBASTIÁN	Santiago	Santiago	51513	Técnico en Logística y similares	TNS en Logística	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2024	5		2024001	Técnico Nivel Sup. en Logística	I162S2C79J4V1	Pregrado	2	79	4	1	I162S2C79J4V1|51513|TNS EN LOGÍSTICA|OTRO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	34525	Ingeniería en Computación e Informática y similares	Ingeniería en Informática (P.E.)	Vespertino	Programa Especial	Profesional	Ingreso Directo	2016	5	2024001001	2024001	Ingeniero (a) en Informática	I162S2C1J2V3	Pregrado	2	1	2	3	I162S2C1J2V3|34525|INGENIERÍA EN INFORMÁTICA (P.E.)|VESPERTINO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40876	Ingeniería en Computación e Informática y similares	Ingeniería en Informática (P.E.)	Otro	Programa Especial	Profesional	Ingreso Directo	2019	4	2024001001	2024001	Ingeniero en Informática	I162S2C1J4V2	Pregrado	2	1	4	2	I162S2C1J4V2|40876|INGENIERÍA EN INFORMÁTICA (P.E.)|OTRO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32512	Ingeniería en Automatización, Control Industrial y similares	Ingeniería en Automatización y Control Industrial	Vespertino	Programa Regular	Profesional	Ingreso Directo	2015	8	2024001001	2024001	Ingeniero  en Automatización y Control Industrial	I162S2C22J2V1	Pregrado	2	22	2	1	I162S2C22J2V1|32512|INGENIERÍA EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|VESPERTINO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6681	Técnico Programador Computacional	TNS en Programación Computacional	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	1990	4		2024001	Técnico Nivel Sup. Programador Computacional	I162S2C2J1V1	Pregrado	2	2	1	1	I162S2C2J1V1|6681|TNS EN PROGRAMACIÓN COMPUTACIONAL|DIURNO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	36742	Técnico Programador Computacional	TNS en Programación Computacional	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2017	5	2024001001	2024001	Técnico Nivel Sup. en Programación Computacional	I162S2C2J4V1	Pregrado	2	2	4	1	I162S2C2J4V1|36742|TNS EN PROGRAMACIÓN COMPUTACIONAL|OTRO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	36743	Técnico en Análisis de Sistemas Informáticos	TNS Análisis de Sistemas	Otro	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2017	6	2024001001	2024001	Técnico Nivel Sup. Análisis de Sistemas	I162S2C35J4V1	Pregrado	2	35	4	1	I162S2C35J4V1|36743|TNS ANÁLISIS DE SISTEMAS|OTRO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	34541	Ingeniería en Conectividad y Redes	Ingeniería en Conectividad y Redes (P.E.)	Vespertino	Programa Especial	Profesional	Ingreso Directo	2016	5	2024001001	2024001	Ingeniero (a) en Conectividad y Redes	I162S2C3J2V3	Pregrado	2	3	2	3	I162S2C3J2V3|34541|INGENIERÍA EN CONECTIVIDAD Y REDES (P.E.)|VESPERTINO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	42539	Ingeniería en Conectividad y Redes	Ingeniería en Conectividad y Redes (P.E.)	Otro	Programa Especial	Profesional	Ingreso Directo	2019	4	2024001001	2024001	Ingeniero en Conectividad y Redes	I162S2C3J4V3	Pregrado	2	3	4	3	I162S2C3J4V3|42539|INGENIERÍA EN CONECTIVIDAD Y REDES (P.E.)|OTRO	OK	REVISAR_COD_CARRERA
2022	2024	IP SAN SEBASTIÁN	Santiago	Santiago	43929	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)	Otro	Programa Especial	Profesional	No es Ingreso Directo	2020	5	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J4V1	Pregrado	2	46	4	1	I162S2C46J4V1|43929|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|OTRO	OK	REVISAR_COD_CARRERA
2021	2024	IP SAN SEBASTIÁN	Santiago	Santiago	38818	Ingeniería en Automatización, Control Industrial y similares	Ingeniero en Automatización y Control Industrial (P.E.)	Vespertino	Programa Especial	Profesional	Ingreso Directo	2018	4	2024001001	2024001	Ingeniero en Automatización y Control Industrial	I162S2C22J2V2	Pregrado	2	22	2	2	I162S2C22J2V2|38818|INGENIERO EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL (P.E.)|VESPERTINO	OK	REVISAR_COD_CARRERA
2021	2024	IP SAN SEBASTIÁN	Santiago	Santiago	6682	Técnico Programador Computacional	TNS en Programación Computacional	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	1990	4		2024001	Técnico Nivel Sup. Programador Computacional	I162S2C2J2V1	Pregrado	2	2	2	1	I162S2C2J2V1|6682|TNS EN PROGRAMACIÓN COMPUTACIONAL|VESPERTINO	OK	REVISAR_COD_CARRERA
2021	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32516	Técnico en Análisis de Sistemas Informáticos	TNS Análisis de Sistemas	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2015	6	2024001001	2024001	Técnico Nivel Sup. Análisis de Sistemas	I162S2C35J2V1	Pregrado	2	35	2	1	I162S2C35J2V1|32516|TNS ANÁLISIS DE SISTEMAS|VESPERTINO	OK	REVISAR_COD_CARRERA
2021	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40715	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)	Vespertino	Programa Especial	Profesional	Ingreso Directo	2019	5	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J2V2	Pregrado	2	46	2	2	I162S2C46J2V2|40715|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|VESPERTINO	OK	REVISAR_COD_CARRERA
2021	2024	IP SAN SEBASTIÁN	Santiago	Santiago	40716	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)	Vespertino	Programa Especial	Profesional	Ingreso Directo	2019	4	2024001001	2024001	Ingeniero en Ciberseguridad	I162S2C46J2V3	Pregrado	2	46	2	3	I162S2C46J2V3|40716|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|VESPERTINO	OK	REVISAR_COD_CARRERA
2019	2024	IP SAN SEBASTIÁN	Santiago	Santiago	28356	Técnico en Prevención de Riesgos	TNS en Prevención de Riesgos	Vespertino	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2013	5	2024001001	2024001	Técnico Nivel Sup. en Prevención de Riesgos	I162S2C10J2V1	Pregrado	2	10	2	1	I162S2C10J2V1|28356|TNS EN PREVENCIÓN DE RIESGOS|VESPERTINO	OK	REVISAR_COD_CARRERA
2019	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32511	Técnico en Automatización, Control automático y similares	TNS en Automatización y Control Industrial	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2015	5	2024001001	2024001	Técnico Nivel Sup. en Automatización y Control Industrial	I162S2C11J1V1	Pregrado	2	11	1	1	I162S2C11J1V1|32511|TNS EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|DIURNO	OK	REVISAR_COD_CARRERA
2019	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32517	Técnico en Análisis de Sistemas Informáticos	TNS en Análisis de Sistemas	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2015	6	2024001001	2024001	Técnico Nivel Sup. en Análisis de Sistemas	I162S2C35J1V1	Pregrado	2	35	1	1	I162S2C35J1V1|32517|TNS EN ANÁLISIS DE SISTEMAS|DIURNO	OK	REVISAR_COD_CARRERA
2015	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32510	Ingeniería en Computación e Informática y similares	Ingeniería en Informática	Diurno	Programa Regular	Profesional	Ingreso Directo	2015	8	2024001001	2024001	Ingeniería en Informática	I162S2C1J1V1	Pregrado	2	1	1	1	I162S2C1J1V1|32510|INGENIERÍA EN INFORMÁTICA|DIURNO	OK	REVISAR_COD_CARRERA
2015	2024	IP SAN SEBASTIÁN	Santiago	Santiago	32514	Ingeniería en Computación e Informática y similares	Ingeniería en Informática	Vespertino	Programa Regular	Profesional	Ingreso Directo	2015	8	2024001001	2024001	Ingeniería en Informática	I162S2C1J2V1	Pregrado	2	1	2	1	I162S2C1J2V1|32514|INGENIERÍA EN INFORMÁTICA|VESPERTINO	OK	REVISAR_COD_CARRERA
2014	2024	IP SAN SEBASTIÁN	Santiago	Santiago	28354	Técnico en Minas y similares	TNS en Minería y Operaciones de Planta	Diurno	Programa Regular	Técnico Nivel Superior	Ingreso Directo	2013	5	2024001001	2024001	Técnico Nivel Sup. en Minería y Operaciones de Planta	I162S2C11J1V1	Pregrado	2	11	1	1	I162S2C11J1V1|28354|TNS EN MINERÍA Y OPERACIONES DE PLANTA|DIURNO	OK	REVISAR_COD_CARRERA
2014	2024	IP SAN SEBASTIÁN	Santiago	Santiago	29353	Ingeniería en Minas	Ing. Ejec. en Minería y Operaciones de Planta	Diurno	Programa Regular	Profesional	Ingreso Directo	2014	8	2024001001	2024001	Ingeniero Ejec. en Minería y Operaciones de Planta	I162S2C22J1V1	Pregrado	2	22	1	1	I162S2C22J1V1|29353|ING. EJEC. EN MINERÍA Y OPERACIONES DE PLANTA|DIURNO	OK	REVISAR_COD_CARRERA
2014	2024	IP SAN SEBASTIÁN	Santiago	Santiago	29354	Ingeniería en Minas	Ing. Ejec. en Minería y Operaciones de Planta	Vespertino	Programa Regular	Profesional	Ingreso Directo	2014	8	2024001001	2024001	Ingeniero Ejec. en Minería y Operaciones de Planta	I162S2C22J2V1	Pregrado	2	22	2	1	I162S2C22J2V1|29354|ING. EJEC. EN MINERÍA Y OPERACIONES DE PLANTA|VESPERTINO	OK	REVISAR_COD_CARRERA
"""

SIN_CODIGO_SIES_TSV = """Año	Cód. Institución	Nombre Institución	Tipo Institución	Clasificación1	Clasificación2	Clasificación3	Clasificación4	Clasificación5	Clasificación6	Nombre de la Sede	Comuna donde se imparte la carrera o programa	Nombre Region	Orden Geográfico de la Región (Norte aSur)	Cód. Carrera	Carrera Genérica	Nombre Programa	Mención o Especialidad	Horario	Tipo Programa	Area Conocimiento	idgenerocarrera	Tipo Carrera	IngresoDirecto	Año Inicio Actividades	Nombre del Campus	Duración (en semestres)	Cód. Campus	Cód. Sede	Título	Grado Académico	Máximo Puntaje (promedio matemáticas y lenguaje)	Promedio Puntaje (promedio matemáticas y lenguaje)	Mínimo Puntaje (promedio matemáticas y lenguaje)	Puntaje de corte (primer seleccionado)	Puntaje de corte (promedio de la carrera)	Puntaje de corte (último seleccionado)	Máximo Puntaje NEM	Promedio Puntaje NEM	Mínimo Puntaje NEM	Máximo Puntaje Ranking	Promedio Puntaje Ranking	Mínimo Puntaje Ranking	Nº Alumnos Ingreso Via PSU o PDT	Nº Alumnos Ingreso Otra Via	Valor de matrícula	Valor de arancel	Valor del Título	Tipo Moneda	Vacantes	Matrícula primer año hombres	Matrícula primer año mujeres	Matrícula primer año extranjeros	Matrícula Primer Año	Matrícula total hombres	Matrícula total mujeres	Matrícula total extranjeros	Matrícula Total	Código SIES	Pregrado/Posgrado	Matrícula Primer Año No Binario	Matrícula Total No Binario	FILA_ORIGINAL	MOTIVO_REVISION
2022	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(b) Subsistema Técnico Profesional	Santiago	Santiago	Región Metropolitana	7	43934	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)		Otro	Programa Especial	Tecnología	1008039	Profesional	No es Ingreso Directo	2020	República	5	2024001001	2024001	Ingeniero en Ciberseguridad															45000	pesos     		0	0	0	0	1	0	0	1		Pregrado	0	0	84	SIN_CODIGO_SIES
2021	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(b) Subsistema Técnico Profesional	Santiago	Santiago	Región Metropolitana	7	43934	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)		Otro	Programa Especial	Tecnología	1008039	Profesional	No es Ingreso Directo	2020	República	5	2024001001	2024001	Ingeniero en Ciberseguridad															55000	Pesos     		0	0	0	0	16	3	0	19		Pregrado	0	0	112	SIN_CODIGO_SIES
2020	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(c) No adscrito	Santiago	Santiago	Región Metropolitana	7	43934	Ingeniería en Computación e Informática y similares	Ingeniería en Ciberseguridad (P.E.)		Otro	Programa Especial	Tecnología	1008039	Profesional	No es Ingreso Directo	2020	República	5	2024001001	2024001	Ingeniero en Ciberseguridad												10	130000	1680000	55000	Pesos     	15	9	1	0	10	9	1	0	10		Pregrado	0	0	143	SIN_CODIGO_SIES
2014	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(c) No adscrito	Santiago	Santiago	Región Metropolitana	7	28353	Técnico en Minas y similares	TNS en Minería y Operaciones de Planta		Vespertino	Programa Regular	Tecnología	1008518	Técnico Nivel Superior	Ingreso Directo	2013	República	5	2024001001	2024001	Técnico Nivel Sup. en Minería y Operaciones de Planta												27	86000	1152000	40000	Pesos     	55	24	3	0	27	46	6	0	52		Pregrado	0	0	275	SIN_CODIGO_SIES
2013	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(c) No adscrito	Santiago	Santiago	Región Metropolitana	7	28353	Técnico en Minas y similares	TNS en Minería y Operaciones de Planta		Vespertino	Programa Regular	Tecnología	1008518	Técnico Nivel Superior	Ingreso Directo	2013	República	5	2024001001	2024001	Técnico Nivel Sup. en Minería y Operaciones de Planta												52	82000	960000	24000	Pesos     	20	44	8	1	52	44	8	1	52		Pregrado	0	0	286	SIN_CODIGO_SIES
2013	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(c) No adscrito	Santiago	Santiago	Región Metropolitana	7	28355	Técnico en Prevención de Riesgos	TNS en Prevención de Riesgos		Diurno	Programa Regular	Tecnología	1008505	Técnico Nivel Superior	Ingreso Directo	2013	República	5	2024001001	2024001	Técnico Nivel Sup. en Prevención de Riesgos												9	82000	960000	24000	Pesos     	20	6	3	1	9	6	3	1	9		Pregrado	0	0	287	SIN_CODIGO_SIES
2020	2024	IP SAN SEBASTIÁN	I.P.	(c) Institutos Profesionales	(e) Institutos Profesionales	(a) Acreditada	(a) Autónoma	(b) No Adscritas/No Aplica	(c) No adscrito	Santiago	Santiago	Región Metropolitana	7	42547	Diplomado	Diplomado en Ciberseguridad Aplicada		Otro	Programa Regular	Tecnología	81501	Diplomado	Ingreso Directo	2019	República	3	2024001001	2024001	Diplomado en Ciberseguridad Aplicada												80000	1200000	0	Pesos     		17	2	1	19	17	2	1	19		Posgrado	0	0	355	SIN_CODIGO_SIES
"""

EXPECTED_YEAR_COUNTS = [
    ("2005", 8),
    ("2006", 8),
    ("2007", 8),
    ("2008", 8),
    ("2009", 8),
    ("2010", 8),
    ("2011", 8),
    ("2012", 8),
    ("2013", 12),
    ("2014", 13),
    ("2015", 13),
    ("2016", 16),
    ("2017", 18),
    ("2018", 22),
    ("2019", 28),
    ("2020", 32),
    ("2021", 31),
    ("2022", 28),
    ("2023", 18),
    ("2024", 25),
    ("2025", 34),
]

EXPECTED_HORARIO = {"Diurno": 14, "Otro": 22, "Vespertino": 23}
EXPECTED_JOR = {"1": 14, "2": 23, "4": 22}


@dataclass(frozen=True)
class Paths:
    repo: Path
    cned: Path
    data: Path
    resultados: Path
    excel_out: Path
    csv_out: Path
    md_out: Path
    script_path: str


def locate_repo() -> Path:
    candidates: list[Path] = [Path.cwd()]
    file_name = globals().get("__file__")
    if file_name:
        try:
            candidates.append(Path(file_name).resolve().parents[3])
        except Exception:
            pass
    for candidate in candidates:
        if (candidate / "indices_2025" / "cned").exists():
            return candidate
    raise FileNotFoundError("No fue posible ubicar el repo avance_curricular.")


def build_paths() -> Paths:
    repo = locate_repo()
    cned = repo / "indices_2025" / "cned"
    data = cned / "data"
    resultados = cned / "resultados"
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    script_name = globals().get("__file__")
    script_path = str(Path(script_name).resolve()) if script_name and Path(script_name).exists() else "stdin"
    return Paths(
        repo=repo,
        cned=cned,
        data=data,
        resultados=resultados,
        excel_out=resultados / f"CNED_ARTEFACTO_OPERATIVO_LIMPIO_{ts}.xlsx",
        csv_out=data / "BASE_CNED_LIMPIA_OPERATIVA.csv",
        md_out=resultados / f"CNED_ARTEFACTO_OPERATIVO_LIMPIO_{ts}.md",
        script_path=script_path,
    )


def clean_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.lower() == "nan":
        return ""
    return text.strip()


def load_embedded_tsv(raw_text: str, right_align_tail: int = 0) -> pd.DataFrame:
    if right_align_tail <= 0:
        df = pd.read_csv(StringIO(raw_text.strip("\n")), sep="\t", dtype=str, keep_default_na=False)
        for col in df.columns:
            df[col] = df[col].map(clean_cell)
        return df

    lines = [line for line in raw_text.strip("\n").splitlines() if line.strip()]
    header = [clean_cell(item) for item in lines[0].split("\t")]
    rows: list[list[str]] = []
    for line in lines[1:]:
        parts = [clean_cell(item) for item in line.split("\t")]
        if len(parts) > len(header):
            raise ValueError(f"Fila con más columnas que el encabezado: {len(parts)} > {len(header)}")
        if len(parts) < len(header):
            if len(parts) < right_align_tail:
                raise ValueError("No se puede alinear por la derecha: la fila es demasiado corta")
            leading = parts[:-right_align_tail]
            trailing = parts[-right_align_tail:]
            gap = len(header) - len(leading) - len(trailing)
            if gap < 0:
                raise ValueError("Alineación inválida al reconstruir la fila")
            parts = leading + ([""] * gap) + trailing
        rows.append(parts)
    return pd.DataFrame(rows, columns=header)


def adjust_widths(ws: Any) -> None:
    for col_idx in range(1, ws.max_column + 1):
        values = [clean_cell(ws.cell(row=row_idx, column=col_idx).value) for row_idx in range(1, ws.max_row + 1)]
        width = min(max(max((len(value) for value in values), default=0) + 2, 12), 80)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def style_sheet(ws: Any, title: str) -> None:
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if title == "BASE_CNED_LIMPIA":
        header_map = {clean_cell(ws.cell(1, col_idx).value): col_idx for col_idx in range(1, ws.max_column + 1)}
        ok_col = header_map.get("VALIDACION_CODIGO_UNICO")
        warn_col = header_map.get("VALIDACION_COD_CARRERA")
        key_col = header_map.get("CLAVE_OPERATIVA")
        for row_idx in range(2, ws.max_row + 1):
            if ok_col:
                ws.cell(row_idx, ok_col).fill = OK_FILL
            if warn_col:
                ws.cell(row_idx, warn_col).fill = WARN_FILL
            if key_col:
                ws.cell(row_idx, key_col).fill = NOTE_FILL
    if title == "SIN_CODIGO_SIES":
        header_map = {clean_cell(ws.cell(1, col_idx).value): col_idx for col_idx in range(1, ws.max_column + 1)}
        motivo_col = header_map.get("MOTIVO_REVISION")
        if motivo_col:
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, motivo_col).fill = WARN_FILL
    if title == "DUPLICADOS_CNED":
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row_idx, col_idx).fill = NOTE_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    adjust_widths(ws)


def add_dataframe_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    for row in dataframe_to_rows(df.fillna(""), index=False, header=True):
        ws.append(row)
    style_sheet(ws, title)


def build_resumen_df(timestamp_label: str) -> pd.DataFrame:
    rows: list[tuple[str, str, str, str]] = [
        ("Fuente", "Hoja1", "", ""),
        ("Generado", timestamp_label, "", ""),
        ("", "", "", ""),
        ("CONTEOS GLOBALES", "", "", ""),
        ("Métrica", "Valor", "Detalle", ""),
        ("Total de filas originales", "354", "Excluye encabezado", ""),
        ("Total de filas con Código SIES vacío", "7", "Enviadas a SIN_CODIGO_SIES", ""),
        ("Total de filas con Código SIES válido", "347", "No vacío; formato se valida aparte", ""),
        ("Total Código SIES con formato OK", "347", "Formato I162S#C#J#V#", ""),
        ("Total Código SIES con formato REVISAR", "0", "No cumple formato esperado", ""),
        ("Total de claves operativas únicas", "59", "Código SIES | Cód. Carrera | Nombre Programa | Horario", ""),
        ("Total de duplicados eliminados", "288", "Registrados en DUPLICADOS_CNED", ""),
        ("Total de registros finales en BASE_CNED_LIMPIA", "59", "Un registro por CLAVE_OPERATIVA", ""),
        ("Años presentes en la base original", "2005, 2006, 2007, 2008, 2009, 2010, 2011, 2012, 2013, 2014, 2015, 2016, 2017, 2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025", "", ""),
        ("Año mínimo", "2005", "Base original", ""),
        ("Año máximo", "2025", "Base original", ""),
        ("", "", "", ""),
        ("CANTIDAD DE REGISTROS POR AÑO - BASE ORIGINAL", "", "", ""),
        ("Año", "Registros", "", ""),
    ]
    rows.extend((year, str(count), "", "") for year, count in EXPECTED_YEAR_COUNTS)
    rows.extend(
        [
            ("", "", "", ""),
            ("CANTIDAD DE REGISTROS POR PREGRADO/POSGRADO - BASE FINAL", "", "", ""),
            ("Pregrado/Posgrado", "Registros", "", ""),
            ("Pregrado", "59", "", ""),
            ("", "", "", ""),
            ("CANTIDAD DE REGISTROS POR HORARIO - BASE FINAL", "", "", ""),
            ("Horario", "Registros", "", ""),
        ]
    )
    rows.extend((label, str(count), "", "") for label, count in EXPECTED_HORARIO.items())
    rows.extend(
        [
            ("", "", "", ""),
            ("CANTIDAD DE REGISTROS POR JOR_DERIVADA - BASE FINAL", "", "", ""),
            ("JOR_DERIVADA", "Registros", "", ""),
        ]
    )
    rows.extend((label, str(count), "", "") for label, count in EXPECTED_JOR.items())
    rows.extend(
        [
            ("", "", "", ""),
            ("VALIDACIONES EN BASE FINAL", "", "", ""),
            ("Validación", "Resultado", "Registros", ""),
            ("VALIDACION_CODIGO_UNICO", "OK", "59", ""),
            ("VALIDACION_COD_CARRERA", "REVISAR_COD_CARRERA", "59", ""),
        ]
    )
    return pd.DataFrame(rows, columns=["RESUMEN_LIMPIEZA", "Columna1", "Columna2", "Columna3"])


def build_duplicates_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "MENSAJE": "Detalle completo de DUPLICADOS_CNED no fue pegado en este prompt",
                "TOTAL_DUPLICADOS_DECLARADOS": "288",
                "CRITERIO_DUPLICIDAD": "Duplicidad operativa por CLAVE_OPERATIVA",
                "OBSERVACION": "Hoja creada para trazabilidad; completar con detalle si se requiere auditoría fila a fila",
            }
        ]
    )


def validate_base_df(df: pd.DataFrame) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    checks: list[str] = []

    if len(df) != 59:
        errors.append(f"BASE_CNED_LIMPIA debe tener 59 filas y tiene {len(df)}")
    else:
        checks.append("BASE_CNED_LIMPIA = 59 filas")

    empty_codes = int(df["Código SIES"].eq("").sum())
    if empty_codes != 0:
        errors.append(f"BASE_CNED_LIMPIA debe tener 0 Código SIES vacíos y tiene {empty_codes}")
    else:
        checks.append("Código SIES vacíos = 0")

    unique_keys = int(df["CLAVE_OPERATIVA"].nunique())
    if unique_keys != 59:
        errors.append(f"CLAVE_OPERATIVA únicas debe ser 59 y es {unique_keys}")
    else:
        checks.append("CLAVE_OPERATIVA únicas = 59")

    ok_unique = int(df["VALIDACION_CODIGO_UNICO"].eq("OK").sum())
    if ok_unique != 59:
        errors.append(f"VALIDACION_CODIGO_UNICO OK debe ser 59 y es {ok_unique}")
    else:
        checks.append("VALIDACION_CODIGO_UNICO = OK en 59 registros")

    warn_count = int(df["VALIDACION_COD_CARRERA"].eq("REVISAR_COD_CARRERA").sum())
    if warn_count != 59:
        errors.append(f"VALIDACION_COD_CARRERA REVISAR_COD_CARRERA debe ser 59 y es {warn_count}")
    else:
        checks.append("VALIDACION_COD_CARRERA = REVISAR_COD_CARRERA en 59 registros")

    pregrado_count = int(df["Pregrado/Posgrado"].eq("Pregrado").sum())
    if pregrado_count != 59:
        errors.append(f"Pregrado/Posgrado debe ser Pregrado en 59 registros y es {pregrado_count}")
    else:
        checks.append("Pregrado/Posgrado = Pregrado en 59 registros")

    horario_counts = {str(key): int(value) for key, value in df["Horario"].value_counts().sort_index().items()}
    if horario_counts != EXPECTED_HORARIO:
        errors.append(f"Distribución Horario esperada {EXPECTED_HORARIO} y obtenida {horario_counts}")
    else:
        checks.append("Distribución Horario = Diurno 14, Otro 22, Vespertino 23")

    jor_counts = {str(key): int(value) for key, value in df["JOR_DERIVADA"].value_counts().sort_index().items()}
    if jor_counts != EXPECTED_JOR:
        errors.append(f"Distribución JOR_DERIVADA esperada {EXPECTED_JOR} y obtenida {jor_counts}")
    else:
        checks.append("Distribución JOR_DERIVADA = 1:14, 2:23, 4:22")

    invalid_format = 0
    derivative_mismatches = 0
    for _, row in df.iterrows():
        match = CODE_PATTERN.fullmatch(row["Código SIES"])
        if not match:
            invalid_format += 1
            continue
        if match.group("sed") != row["COD_SED_DERIVADO"]:
            derivative_mismatches += 1
        if match.group("car") != row["COD_CAR_DERIVADO"]:
            derivative_mismatches += 1
        if match.group("jor") != row["JOR_DERIVADA"]:
            derivative_mismatches += 1
        if match.group("version") != row["VERSION_DERIVADA"]:
            derivative_mismatches += 1

    if invalid_format != 0:
        errors.append(f"Código SIES con formato inválido: {invalid_format}")
    else:
        checks.append("Formato Código SIES = OK en 59 registros")

    if derivative_mismatches != 0:
        errors.append(f"Derivados del Código SIES inconsistentes: {derivative_mismatches}")
    else:
        checks.append("Derivados del Código SIES coherentes con COD_SED/COD_CAR/JOR/VERSION")

    return checks, errors


def validate_sin_df(df: pd.DataFrame) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    checks: list[str] = []
    if len(df) != 7:
        errors.append(f"SIN_CODIGO_SIES debe tener 7 filas y tiene {len(df)}")
    else:
        checks.append("SIN_CODIGO_SIES = 7 filas")
    empty_codes = int(df["Código SIES"].eq("").sum())
    if empty_codes != 7:
        errors.append(f"SIN_CODIGO_SIES debe tener 7 Código SIES vacíos y tiene {empty_codes}")
    else:
        checks.append("SIN_CODIGO_SIES conserva 7 Código SIES vacíos")
    return checks, errors


def build_markdown(paths: Paths, base_df: pd.DataFrame, sin_df: pd.DataFrame, checks: list[str]) -> str:
    warnings = [
        "REVISAR_COD_CARRERA es advertencia no bloqueante; Cód. Carrera CNED y COD_CAR_DERIVADO no son equivalentes directos.",
        "DUPLICADOS_CNED no incluye el detalle fila a fila porque ese detalle no fue pegado en el prompt; se registran los 288 duplicados declarados para trazabilidad.",
    ]
    rules = [
        "No modificar Hoja1.",
        "No modificar archivos originales.",
        "No actualizar DATOS_VALIDACION_CNED_TSV.tsv.",
        "No actualizar DATOS_VALIDACION_CNED_TSV_ACTUALIZADO.tsv.",
        "Usar BASE_CNED_LIMPIA en grano CLAVE_OPERATIVA.",
        "Validar operativamente con Código SIES, COD_SED_DERIVADO, COD_CAR_DERIVADO, JOR_DERIVADA, VERSION_DERIVADA y CLAVE_OPERATIVA.",
        "Mantener SIN_CODIGO_SIES separado para revisión manual.",
    ]
    checks_md = "\n".join(f"- {item}" for item in checks)
    warnings_md = "\n".join(f"- {item}" for item in warnings)
    rules_md = "\n".join(f"- {item}" for item in rules)
    return f"""# CNED Artefacto Operativo Limpio

## Dictamen

{DICTAMEN}

## Archivos generados

- Excel: {paths.excel_out.resolve()}
- CSV: {paths.csv_out.resolve()}
- Markdown: {paths.md_out.resolve()}
- Script: {paths.script_path}

## Conteos

- BASE_CNED_LIMPIA: {len(base_df)} registros
- SIN_CODIGO_SIES: {len(sin_df)} registros
- DUPLICADOS_CNED declarados: 288 registros
- Código SIES vacíos en BASE_CNED_LIMPIA: 0
- CLAVE_OPERATIVA únicas en BASE_CNED_LIMPIA: {base_df['CLAVE_OPERATIVA'].nunique()}
- VALIDACION_CODIGO_UNICO = OK: {int(base_df['VALIDACION_CODIGO_UNICO'].eq('OK').sum())}
- VALIDACION_COD_CARRERA = REVISAR_COD_CARRERA: {int(base_df['VALIDACION_COD_CARRERA'].eq('REVISAR_COD_CARRERA').sum())}

## Validaciones principales

{checks_md}

## Advertencias

{warnings_md}

## Reglas de uso

{rules_md}

## Cierre

Con base exclusivamente en la información entregada en el prompt, el dictamen es BASE_CNED_LIMPIA_LISTA_PARA_USO_OPERATIVO. Sin embargo, el artefacto operativo aún debía materializarse en el proyecto, ya que el TSV manual disponible contenía solo encabezados y 0 filas.

Este proceso materializa un artefacto operativo controlado en Excel y CSV, sin modificar la hoja original, sin actualizar el TSV base y sin automatizar sustituciones. La base final queda en el grano CLAVE_OPERATIVA, con 59 registros de Pregrado, 0 códigos SIES vacíos, 0 claves duplicadas y validación fuerte sobre Código SIES y sus derivados.

La advertencia REVISAR_COD_CARRERA no bloquea el uso de la base, porque Cód. Carrera CNED y COD_CAR_DERIVADO pertenecen a naturalezas distintas. La validación operativa debe realizarse sobre Código SIES, COD_SED_DERIVADO, COD_CAR_DERIVADO, JOR_DERIVADA, VERSION_DERIVADA y CLAVE_OPERATIVA.
"""


def main() -> int:
    paths = build_paths()
    paths.data.mkdir(parents=True, exist_ok=True)
    paths.resultados.mkdir(parents=True, exist_ok=True)

    base_df = load_embedded_tsv(BASE_CNED_LIMPIA_TSV)
    sin_df = load_embedded_tsv(SIN_CODIGO_SIES_TSV, right_align_tail=8)
    duplicates_df = build_duplicates_df()
    resumen_df = build_resumen_df(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

    base_checks, base_errors = validate_base_df(base_df)
    sin_checks, sin_errors = validate_sin_df(sin_df)
    checks = [*base_checks, *sin_checks]
    errors = [*base_errors, *sin_errors]

    if errors:
        print("=" * 120)
        print("DICTAMEN: ERROR_EN_VALIDACIONES")
        for error in errors:
            print(f"- {error}")
        print("=" * 120)
        return 1

    wb = Workbook()
    wb.remove(wb.active)
    for title, df in [
        ("RESUMEN_LIMPIEZA", resumen_df),
        ("BASE_CNED_LIMPIA", base_df),
        ("SIN_CODIGO_SIES", sin_df),
        ("DUPLICADOS_CNED", duplicates_df),
    ]:
        add_dataframe_sheet(wb, title, df)
    wb.save(paths.excel_out)

    base_df.to_csv(paths.csv_out, sep=";", index=False, encoding="utf-8")
    markdown = build_markdown(paths, base_df, sin_df, checks)
    paths.md_out.write_text(markdown, encoding="utf-8")

    print("=" * 120)
    print("CNED — Materialización de artefacto operativo limpio")
    print("=" * 120)
    print(f"Excel generado: {paths.excel_out.resolve()}")
    print(f"CSV generado: {paths.csv_out.resolve()}")
    print(f"Markdown generado: {paths.md_out.resolve()}")
    print(f"Dictamen final: {DICTAMEN}")
    print("Validaciones principales:")
    for item in checks:
        print(f"- {item}")
    print("Advertencias:")
    print("- REVISAR_COD_CARRERA queda como advertencia no bloqueante.")
    print("- DUPLICADOS_CNED queda como hoja de trazabilidad sin detalle fila a fila; total declarado 288.")
    print("=" * 120)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
r'''
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Materializa un artefacto operativo CNED limpio y trazable."""

from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from io import StringIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows


DICTAMEN = "BASE_CNED_LIMPIA_LISTA_PARA_USO_OPERATIVO"
EXPECTED_BASE_ROWS = 59
EXPECTED_SIN_ROWS = 7
EXPECTED_BASE_COLS = 26
EXPECTED_SIN_COLS = 64
EXPECTED_HORARIO = {"Diurno": 14, "Otro": 22, "Vespertino": 23}
EXPECTED_JOR = {"1": 14, "2": 23, "4": 22}
EXPECTED_PREGRADO = {"Pregrado": 59}
EXPECTED_VALIDACION_CODIGO = {"OK": 59}
EXPECTED_VALIDACION_CARRERA = {"REVISAR_COD_CARRERA": 59}
EXPECTED_SIN_PREGRADO = {"Pregrado": 6, "Posgrado": 1}
EXPECTED_SIN_MOTIVO = {"SIN_CODIGO_SIES": 7}
EXPECTED_DUPLICADOS = 288

HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="E2F0D9")
OK_FILL = PatternFill("solid", fgColor="E2F0D9")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FILL = PatternFill("solid", fgColor="FCE4D6")

BASE_CNED_LIMPIA_TSV = (
    "Año\tCód. Institución\tNombre Institución\tNombre de la Sede\tComuna donde se imparte la carrera o programa\tCód. Carrera\tCarrera Genérica\tNombre Programa\tHorario\tTipo Programa\tTipo Carrera\tIngresoDirecto\tAño Inicio Actividades\tDuración (en semestres)\tCód. Campus\tCód. Sede\tTítulo\tCódigo SIES\tPregrado/Posgrado\tCOD_SED_DERIVADO\tCOD_CAR_DERIVADO\tJOR_DERIVADA\tVERSION_DERIVADA\tCLAVE_OPERATIVA\tVALIDACION_CODIGO_UNICO\tVALIDACION_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32513\tTécnico en Automatización, Control automático y similares\tTNS en Automatización y Control Industrial\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2015\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Automatización y Control Industrial\tI162S2C11J2V1\tPregrado\t2\t11\t2\t1\tI162S2C11J2V1|32513|TNS EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6677\tIngeniería en Computación e Informática y similares\tIngeniería en Informática\tDiurno\tPrograma Regular\tProfesional\tIngreso Directo\t1990\t8\t\t2024001\tIngeniero en Informática\tI162S2C1J1V2\tPregrado\t2\t1\t1\t2\tI162S2C1J1V2|6677|INGENIERÍA EN INFORMÁTICA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6678\tIngeniería en Computación e Informática y similares\tIngeniería en Informática\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t1990\t8\t\t2024001\tIngeniero en Informática\tI162S2C1J2V4\tPregrado\t2\t1\t2\t4\tI162S2C1J2V4|6678|INGENIERÍA EN INFORMÁTICA|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6680\tIngeniería en Conectividad y Redes\tIngeniería en Conectividad y Redes\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2002\t8\t\t2024001\tIngeniero en Conectividad y Redes\tI162S2C3J2V4\tPregrado\t2\t3\t2\t4\tI162S2C3J2V4|6680|INGENIERÍA EN CONECTIVIDAD Y REDES|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t38816\tIngeniería en Conectividad y Redes\tIngeniería en Conectividad y Redes\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2018\t8\t2024001001\t2024001\tIngeniero en Conectividad y Redes\tI162S2C3J4V2\tPregrado\t2\t3\t4\t2\tI162S2C3J4V2|38816|INGENIERÍA EN CONECTIVIDAD Y REDES|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40713\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad\tDiurno\tPrograma Regular\tProfesional\tIngreso Directo\t2019\t8\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J1V1\tPregrado\t2\t46\t1\t1\tI162S2C46J1V1|40713|INGENIERÍA EN CIBERSEGURIDAD|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40714\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2019\t8\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J2V1\tPregrado\t2\t46\t2\t1\tI162S2C46J2V1|40714|INGENIERÍA EN CIBERSEGURIDAD|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t38817\tIngeniería en Computación e Informática y similares\tIngeniería en Informática\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2018\t8\t2024001001\t2024001\tIngeniero en Informática\tI162S2C46J4V2\tPregrado\t2\t46\t4\t2\tI162S2C46J4V2|38817|INGENIERÍA EN INFORMÁTICA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43928\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2020\t4\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J4V3\tPregrado\t2\t46\t4\t3\tI162S2C46J4V3|43928|INGENIERÍA EN CIBERSEGURIDAD|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40717\tTécnico en Análisis de Sistemas Informáticos\tTNS en Ciberseguridad\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2019\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Ciberseguridad\tI162S2C47J1V1\tPregrado\t2\t47\t1\t1\tI162S2C47J1V1|40717|TNS EN CIBERSEGURIDAD|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40711\tTécnico en Análisis de Sistemas Informáticos\tTNS en Ciberseguridad\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2019\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Ciberseguridad\tI162S2C47J2V1\tPregrado\t2\t47\t2\t1\tI162S2C47J2V1|40711|TNS EN CIBERSEGURIDAD|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43931\tTécnico en Análisis de Sistemas Informáticos\tTNS en Ciberseguridad\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2020\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Ciberseguridad\tI162S2C47J4V1\tPregrado\t2\t47\t4\t1\tI162S2C47J4V1|43931|TNS EN CIBERSEGURIDAD|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43935\tTécnico Programador Computacional\tTNS en Programación y Análisis de Sistemas\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2020\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Programación y Análisis de Sistemas\tI162S2C57J1V1\tPregrado\t2\t57\t1\t1\tI162S2C57J1V1|43935|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43930\tTécnico Programador Computacional\tTNS en Programación y Análisis de Sistemas\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2020\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Programación y Análisis de Sistemas\tI162S2C57J2V1\tPregrado\t2\t57\t2\t1\tI162S2C57J2V1|43930|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43932\tTécnico Programador Computacional\tTNS en Programación y Análisis de Sistemas\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2020\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Programación Computacional y Análisis de Sistemas\tI162S2C57J4V1\tPregrado\t2\t57\t4\t1\tI162S2C57J4V1|43932|TNS EN PROGRAMACIÓN Y ANÁLISIS DE SISTEMAS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6684\tTécnico en Administración de Redes Computacionales\tTNS en Conectividad y Redes\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2002\t4\t\t2024001\tTécnico Nivel Sup. en Conectividad y Redes\tI162S2C6J2V2\tPregrado\t2\t6\t2\t2\tI162S2C6J2V2|6684|TNS EN CONECTIVIDAD Y REDES|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t36744\tTécnico en Administración de Redes Computacionales\tTNS en Conectividad y Redes\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2017\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Conectividad y Redes\tI162S2C6J4V2\tPregrado\t2\t6\t4\t2\tI162S2C6J4V2|36744|TNS EN CONECTIVIDAD Y REDES|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51504\tIngeniería en Administración, Administración de Empresas y similares\tIngeniería en Administración de Empresas\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2024\t8\t\t2024001\tIngeniero en Administración de Empresas\tI162S2C76J2V1\tPregrado\t2\t76\t2\t1\tI162S2C76J2V1|51504|INGENIERÍA EN ADMINISTRACIÓN DE EMPRESAS|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51505\tIngeniería en Administración, Administración de Empresas y similares\tIngeniería en Administración de Empresas\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2024\t8\t\t2024001\tIngeniero en Administración de Empresas\tI162S2C76J4V1\tPregrado\t2\t76\t4\t1\tI162S2C76J4V1|51505|INGENIERÍA EN ADMINISTRACIÓN DE EMPRESAS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51508\tIngeniería en Logística y similares\tIngeniería en Logística\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2024\t8\t\t2024001\tIngeniero en Logística\tI162S2C77J2V1\tPregrado\t2\t77\t2\t1\tI162S2C77J2V1|51508|INGENIERÍA EN LOGÍSTICA|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51509\tIngeniería en Logística y similares\tIngeniería en Logística\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2024\t8\t\t2024001\tIngeniero en Logística\tI162S2C77J4V1\tPregrado\t2\t77\t4\t1\tI162S2C77J4V1|51509|INGENIERÍA EN LOGÍSTICA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51506\tTécnico en Administración de Empresas\tTNS en Administración de Empresas\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2024\t5\t\t2024001\tTécnico Nivel Sup. en Administración de Empresas\tI162S2C78J2V1\tPregrado\t2\t78\t2\t1\tI162S2C78J2V1|51506|TNS EN ADMINISTRACIÓN DE EMPRESAS|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51510\tTécnico en Administración de Empresas\tTNS en Administración de Empresas\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2024\t5\t\t2024001\tTécnico Nivel Sup. en Administración de Empresas\tI162S2C78J4V1\tPregrado\t2\t78\t4\t1\tI162S2C78J4V1|51510|TNS EN ADMINISTRACIÓN DE EMPRESAS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51511\tTécnico en Logística y similares\tTNS en Logística\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2024\t5\t\t2024001\tTécnico Nivel Sup. en Logística\tI162S2C79J1V1\tPregrado\t2\t79\t1\t1\tI162S2C79J1V1|51511|TNS EN LOGÍSTICA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51512\tTécnico en Logística y similares\tTNS en Logística\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2024\t5\t\t2024001\tTécnico Nivel Sup. en Logística\tI162S2C79J2V1\tPregrado\t2\t79\t2\t1\tI162S2C79J2V1|51512|TNS EN LOGÍSTICA|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53717\tAdministración pública y similares\tAdministración Pública\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2025\t8\t2024001001\t2024001\tAdministrador Público\tI162S2C83J4V1\tPregrado\t2\t83\t4\t1\tI162S2C83J4V1|53717|ADMINISTRACIÓN PÚBLICA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53720\tTécnico en Administración Pública\tTNS en Administración Pública\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2025\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Administración Pública\tI162S2C84J4V1\tPregrado\t2\t84\t4\t1\tI162S2C84J4V1|53720|TNS EN ADMINISTRACIÓN PÚBLICA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53718\tTécnico en Contabilidad, Auditoría y similares\tTNS en Contabilidad General\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2025\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Contabilidad General\tI162S2C85J4V1\tPregrado\t2\t85\t4\t1\tI162S2C85J4V1|53718|TNS EN CONTABILIDAD GENERAL|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53721\tContabilidad, Auditoría y similares\tAuditoria\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2025\t8\t2024001001\t2024001\tContador Auditor\tI162S2C86J4V1\tPregrado\t2\t86\t4\t1\tI162S2C86J4V1|53721|AUDITORIA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53723\tIngeniería Ejecución Industrial y similares\tIngeniería Industrial (P.E.)\tOtro\tPrograma Especial\tProfesional\tNo es Ingreso Directo\t2025\t4\t2024001001\t2024001\tIngeniero Industrial\tI162S2C87J4V1\tPregrado\t2\t87\t4\t1\tI162S2C87J4V1|53723|INGENIERÍA INDUSTRIAL (P.E.)|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53722\tIngeniería en Computación e Informática y similares\tIngeniería en Ciencia de Datos\tOtro\tPrograma Regular\tProfesional\tIngreso Directo\t2025\t8\t2024001001\t2024001\tIngeniero en Ciencia de Datos\tI162S2C88J4V1\tPregrado\t2\t88\t4\t1\tI162S2C88J4V1|53722|INGENIERÍA EN CIENCIA DE DATOS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53719\tTécnico en Computación e Informática\tTNS en Ciencia de Datos\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2025\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Ciencia de Datos\tI162S2C89J4V1\tPregrado\t2\t89\t4\t1\tI162S2C89J4V1|53719|TNS EN CIENCIA DE DATOS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53725\tTécnico en Enfermería y Similares\tTNS en Enfermería\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2025\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Enfermería\tI162S2C91J1V1\tPregrado\t2\t91\t1\t1\tI162S2C91J1V1|53725|TNS EN ENFERMERÍA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2025\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t53726\tTécnico en Enfermería y Similares\tTNS en Enfermería\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2025\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Enfermería\tI162S2C91J2V1\tPregrado\t2\t91\t2\t1\tI162S2C91J2V1|53726|TNS EN ENFERMERÍA|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2024\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6679\tIngeniería en Conectividad y Redes\tIngeniería en Conectividad y Redes\tDiurno\tPrograma Regular\tProfesional\tIngreso Directo\t2002\t8\t\t2024001\tIngeniería en Conectividad y Redes\tI162S2C3J1V2\tPregrado\t2\t3\t1\t2\tI162S2C3J1V2|6679|INGENIERÍA EN CONECTIVIDAD Y REDES|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2024\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6683\tTécnico en Administración de Redes Computacionales\tTNS en Conectividad y Redes\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2002\t4\t\t2024001\tTécnico Nivel Sup. en Conectividad y Redes\tI162S2C6J1V2\tPregrado\t2\t6\t1\t2\tI162S2C6J1V2|6683|TNS EN CONECTIVIDAD Y REDES|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2024\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t51513\tTécnico en Logística y similares\tTNS en Logística\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2024\t5\t\t2024001\tTécnico Nivel Sup. en Logística\tI162S2C79J4V1\tPregrado\t2\t79\t4\t1\tI162S2C79J4V1|51513|TNS EN LOGÍSTICA|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t34525\tIngeniería en Computación e Informática y similares\tIngeniería en Informática (P.E.)\tVespertino\tPrograma Especial\tProfesional\tIngreso Directo\t2016\t5\t2024001001\t2024001\tIngeniero (a) en Informática\tI162S2C1J2V3\tPregrado\t2\t1\t2\t3\tI162S2C1J2V3|34525|INGENIERÍA EN INFORMÁTICA (P.E.)|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40876\tIngeniería en Computación e Informática y similares\tIngeniería en Informática (P.E.)\tOtro\tPrograma Especial\tProfesional\tIngreso Directo\t2019\t4\t2024001001\t2024001\tIngeniero en Informática\tI162S2C1J4V2\tPregrado\t2\t1\t4\t2\tI162S2C1J4V2|40876|INGENIERÍA EN INFORMÁTICA (P.E.)|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32512\tIngeniería en Automatización, Control Industrial y similares\tIngeniería en Automatización y Control Industrial\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2015\t8\t2024001001\t2024001\tIngeniero en Automatización y Control Industrial\tI162S2C22J2V1\tPregrado\t2\t22\t2\t1\tI162S2C22J2V1|32512|INGENIERÍA EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6681\tTécnico Programador Computacional\tTNS en Programación Computacional\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t1990\t4\t\t2024001\tTécnico Nivel Sup. Programador Computacional\tI162S2C2J1V1\tPregrado\t2\t2\t1\t1\tI162S2C2J1V1|6681|TNS EN PROGRAMACIÓN COMPUTACIONAL|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t36742\tTécnico Programador Computacional\tTNS en Programación Computacional\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2017\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Programación Computacional\tI162S2C2J4V1\tPregrado\t2\t2\t4\t1\tI162S2C2J4V1|36742|TNS EN PROGRAMACIÓN COMPUTACIONAL|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t36743\tTécnico en Análisis de Sistemas Informáticos\tTNS Análisis de Sistemas\tOtro\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2017\t6\t2024001001\t2024001\tTécnico Nivel Sup. Análisis de Sistemas\tI162S2C35J4V1\tPregrado\t2\t35\t4\t1\tI162S2C35J4V1|36743|TNS ANÁLISIS DE SISTEMAS|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t34541\tIngeniería en Conectividad y Redes\tIngeniería en Conectividad y Redes (P.E.)\tVespertino\tPrograma Especial\tProfesional\tIngreso Directo\t2016\t5\t2024001001\t2024001\tIngeniero (a) en Conectividad y Redes\tI162S2C3J2V3\tPregrado\t2\t3\t2\t3\tI162S2C3J2V3|34541|INGENIERÍA EN CONECTIVIDAD Y REDES (P.E.)|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t42539\tIngeniería en Conectividad y Redes\tIngeniería en Conectividad y Redes (P.E.)\tOtro\tPrograma Especial\tProfesional\tIngreso Directo\t2019\t4\t2024001001\t2024001\tIngeniero en Conectividad y Redes\tI162S2C3J4V3\tPregrado\t2\t3\t4\t3\tI162S2C3J4V3|42539|INGENIERÍA EN CONECTIVIDAD Y REDES (P.E.)|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t43929\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\tOtro\tPrograma Especial\tProfesional\tNo es Ingreso Directo\t2020\t5\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J4V1\tPregrado\t2\t46\t4\t1\tI162S2C46J4V1|43929|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|OTRO\tOK\tREVISAR_COD_CARRERA\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t38818\tIngeniería en Automatización, Control Industrial y similares\tIngeniero en Automatización y Control Industrial (P.E.)\tVespertino\tPrograma Especial\tProfesional\tIngreso Directo\t2018\t4\t2024001001\t2024001\tIngeniero en Automatización y Control Industrial\tI162S2C22J2V2\tPregrado\t2\t22\t2\t2\tI162S2C22J2V2|38818|INGENIERO EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL (P.E.)|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t6682\tTécnico Programador Computacional\tTNS en Programación Computacional\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t1990\t4\t\t2024001\tTécnico Nivel Sup. Programador Computacional\tI162S2C2J2V1\tPregrado\t2\t2\t2\t1\tI162S2C2J2V1|6682|TNS EN PROGRAMACIÓN COMPUTACIONAL|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32516\tTécnico en Análisis de Sistemas Informáticos\tTNS Análisis de Sistemas\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2015\t6\t2024001001\t2024001\tTécnico Nivel Sup. Análisis de Sistemas\tI162S2C35J2V1\tPregrado\t2\t35\t2\t1\tI162S2C35J2V1|32516|TNS ANÁLISIS DE SISTEMAS|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40715\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\tVespertino\tPrograma Especial\tProfesional\tIngreso Directo\t2019\t5\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J2V2\tPregrado\t2\t46\t2\t2\tI162S2C46J2V2|40715|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t40716\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\tVespertino\tPrograma Especial\tProfesional\tIngreso Directo\t2019\t4\t2024001001\t2024001\tIngeniero en Ciberseguridad\tI162S2C46J2V3\tPregrado\t2\t46\t2\t3\tI162S2C46J2V3|40716|INGENIERÍA EN CIBERSEGURIDAD (P.E.)|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2019\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t28356\tTécnico en Prevención de Riesgos\tTNS en Prevención de Riesgos\tVespertino\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2013\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Prevención de Riesgos\tI162S2C10J2V1\tPregrado\t2\t10\t2\t1\tI162S2C10J2V1|28356|TNS EN PREVENCIÓN DE RIESGOS|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2019\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32511\tTécnico en Automatización, Control automático y similares\tTNS en Automatización y Control Industrial\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2015\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Automatización y Control Industrial\tI162S2C11J1V1\tPregrado\t2\t11\t1\t1\tI162S2C11J1V1|32511|TNS EN AUTOMATIZACIÓN Y CONTROL INDUSTRIAL|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2019\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32517\tTécnico en Análisis de Sistemas Informáticos\tTNS en Análisis de Sistemas\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2015\t6\t2024001001\t2024001\tTécnico Nivel Sup. en Análisis de Sistemas\tI162S2C35J1V1\tPregrado\t2\t35\t1\t1\tI162S2C35J1V1|32517|TNS EN ANÁLISIS DE SISTEMAS|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2015\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32510\tIngeniería en Computación e Informática y similares\tIngeniería en Informática\tDiurno\tPrograma Regular\tProfesional\tIngreso Directo\t2015\t8\t2024001001\t2024001\tIngeniería en Informática\tI162S2C1J1V1\tPregrado\t2\t1\t1\t1\tI162S2C1J1V1|32510|INGENIERÍA EN INFORMÁTICA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2015\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t32514\tIngeniería en Computación e Informática y similares\tIngeniería en Informática\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2015\t8\t2024001001\t2024001\tIngeniería en Informática\tI162S2C1J2V1\tPregrado\t2\t1\t2\t1\tI162S2C1J2V1|32514|INGENIERÍA EN INFORMÁTICA|VESPERTINO\tOK\tREVISAR_COD_CARRERA\n"
    "2014\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t28354\tTécnico en Minas y similares\tTNS en Minería y Operaciones de Planta\tDiurno\tPrograma Regular\tTécnico Nivel Superior\tIngreso Directo\t2013\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Minería y Operaciones de Planta\tI162S2C11J1V1\tPregrado\t2\t11\t1\t1\tI162S2C11J1V1|28354|TNS EN MINERÍA Y OPERACIONES DE PLANTA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2014\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t29353\tIngeniería en Minas\tIng. Ejec. en Minería y Operaciones de Planta\tDiurno\tPrograma Regular\tProfesional\tIngreso Directo\t2014\t8\t2024001001\t2024001\tIngeniero Ejec. en Minería y Operaciones de Planta\tI162S2C22J1V1\tPregrado\t2\t22\t1\t1\tI162S2C22J1V1|29353|ING. EJEC. EN MINERÍA Y OPERACIONES DE PLANTA|DIURNO\tOK\tREVISAR_COD_CARRERA\n"
    "2014\t2024\tIP SAN SEBASTIÁN\tSantiago\tSantiago\t29354\tIngeniería en Minas\tIng. Ejec. en Minería y Operaciones de Planta\tVespertino\tPrograma Regular\tProfesional\tIngreso Directo\t2014\t8\t2024001001\t2024001\tIngeniero Ejec. en Minería y Operaciones de Planta\tI162S2C22J2V1\tPregrado\t2\t22\t2\t1\tI162S2C22J2V1|29354|ING. EJEC. EN MINERÍA Y OPERACIONES DE PLANTA|VESPERTINO\tOK\tREVISAR_COD_CARRERA"
)

SIN_CODIGO_SIES_TSV = (
    "Año\tCód. Institución\tNombre Institución\tTipo Institución\tClasificación1\tClasificación2\tClasificación3\tClasificación4\tClasificación5\tClasificación6\tNombre de la Sede\tComuna donde se imparte la carrera o programa\tNombre Region\tOrden Geográfico de la Región (Norte aSur)\tCód. Carrera\tCarrera Genérica\tNombre Programa\tMención o Especialidad\tHorario\tTipo Programa\tArea Conocimiento\tidgenerocarrera\tTipo Carrera\tIngresoDirecto\tAño Inicio Actividades\tNombre del Campus\tDuración (en semestres)\tCód. Campus\tCód. Sede\tTítulo\tGrado Académico\tMáximo Puntaje (promedio matemáticas y lenguaje)\tPromedio Puntaje (promedio matemáticas y lenguaje)\tMínimo Puntaje (promedio matemáticas y lenguaje)\tPuntaje de corte (primer seleccionado)\tPuntaje de corte (promedio de la carrera)\tPuntaje de corte (último seleccionado)\tMáximo Puntaje NEM\tPromedio Puntaje NEM\tMínimo Puntaje NEM\tMáximo Puntaje Ranking\tPromedio Puntaje Ranking\tMínimo Puntaje Ranking\tNº Alumnos Ingreso Via PSU o PDT\tNº Alumnos Ingreso Otra Via\tValor de matrícula\tValor de arancel\tValor del Título\tTipo Moneda\tVacantes\tMatrícula primer año hombres\tMatrícula primer año mujeres\tMatrícula primer año extranjeros\tMatrícula Primer Año\tMatrícula total hombres\tMatrícula total mujeres\tMatrícula total extranjeros\tMatrícula Total\tCódigo SIES\tPregrado/Posgrado\tMatrícula Primer Año No Binario\tMatrícula Total No Binario\tFILA_ORIGINAL\tMOTIVO_REVISION\n"
    "2022\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(b) Subsistema Técnico Profesional\tSantiago\tSantiago\tRegión Metropolitana\t7\t43934\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\t\tOtro\tPrograma Especial\tTecnología\t1008039\tProfesional\tNo es Ingreso Directo\t2020\tRepública\t5\t2024001001\t2024001\tIngeniero en Ciberseguridad\t\t\t\t\t\t\t\t\t\t\t\t\t\t\t45000\tpesos\t\t0\t0\t0\t0\t1\t0\t0\t1\t\tPregrado\t0\t0\t84\tSIN_CODIGO_SIES\n"
    "2021\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(b) Subsistema Técnico Profesional\tSantiago\tSantiago\tRegión Metropolitana\t7\t43934\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\t\tOtro\tPrograma Especial\tTecnología\t1008039\tProfesional\tNo es Ingreso Directo\t2020\tRepública\t5\t2024001001\t2024001\tIngeniero en Ciberseguridad\t\t\t\t\t\t\t\t\t\t\t\t\t\t\t55000\tPesos\t\t0\t0\t0\t0\t16\t3\t0\t19\t\tPregrado\t0\t0\t112\tSIN_CODIGO_SIES\n"
    "2020\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(c) No adscrito\tSantiago\tSantiago\tRegión Metropolitana\t7\t43934\tIngeniería en Computación e Informática y similares\tIngeniería en Ciberseguridad (P.E.)\t\tOtro\tPrograma Especial\tTecnología\t1008039\tProfesional\tNo es Ingreso Directo\t2020\tRepública\t5\t2024001001\t2024001\tIngeniero en Ciberseguridad\t\t\t\t\t\t\t\t\t\t\t\t\t10\t130000\t1680000\t55000\tPesos\t15\t9\t1\t0\t10\t9\t1\t0\t10\t\tPregrado\t0\t0\t143\tSIN_CODIGO_SIES\n"
    "2014\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(c) No adscrito\tSantiago\tSantiago\tRegión Metropolitana\t7\t28353\tTécnico en Minas y similares\tTNS en Minería y Operaciones de Planta\t\tVespertino\tPrograma Regular\tTecnología\t1008518\tTécnico Nivel Superior\tIngreso Directo\t2013\tRepública\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Minería y Operaciones de Planta\t\t\t\t\t\t\t\t\t\t\t\t\t27\t86000\t1152000\t40000\tPesos\t55\t24\t3\t0\t27\t46\t6\t0\t52\t\tPregrado\t0\t0\t275\tSIN_CODIGO_SIES\n"
    "2013\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(c) No adscrito\tSantiago\tSantiago\tRegión Metropolitana\t7\t28353\tTécnico en Minas y similares\tTNS en Minería y Operaciones de Planta\t\tVespertino\tPrograma Regular\tTecnología\t1008518\tTécnico Nivel Superior\tIngreso Directo\t2013\tRepública\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Minería y Operaciones de Planta\t\t\t\t\t\t\t\t\t\t\t\t\t52\t82000\t960000\t24000\tPesos\t20\t44\t8\t1\t52\t44\t8\t1\t52\t\tPregrado\t0\t0\t286\tSIN_CODIGO_SIES\n"
    "2013\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(c) No adscrito\tSantiago\tSantiago\tRegión Metropolitana\t7\t28355\tTécnico en Prevención de Riesgos\tTNS en Prevención de Riesgos\t\tDiurno\tPrograma Regular\tTecnología\t1008505\tTécnico Nivel Superior\tIngreso Directo\t2013\tRepública\t5\t2024001001\t2024001\tTécnico Nivel Sup. en Prevención de Riesgos\t\t\t\t\t\t\t\t\t\t\t\t\t9\t82000\t960000\t24000\tPesos\t20\t6\t3\t1\t9\t6\t3\t1\t9\t\tPregrado\t0\t0\t287\tSIN_CODIGO_SIES\n"
    "2020\t2024\tIP SAN SEBASTIÁN\tI.P.\t(c) Institutos Profesionales\t(e) Institutos Profesionales\t(a) Acreditada\t(a) Autónoma\t(b) No Adscritas/No Aplica\t(c) No adscrito\tSantiago\tSantiago\tRegión Metropolitana\t7\t42547\tDiplomado\tDiplomado en Ciberseguridad Aplicada\t\tOtro\tPrograma Regular\tTecnología\t81501\tDiplomado\tIngreso Directo\t2019\tRepública\t3\t2024001001\t2024001\tDiplomado en Ciberseguridad Aplicada\t\t\t\t\t\t\t\t\t\t\t\t\t80000\t1200000\t0\tPesos\t\t17\t2\t1\t19\t17\t2\t1\t19\t\tPosgrado\t0\t0\t355\tSIN_CODIGO_SIES"
)

CIERRE_TEXTO = (
    "Con base exclusivamente en la información entregada en el prompt, el dictamen es "
    "BASE_CNED_LIMPIA_LISTA_PARA_USO_OPERATIVO. Sin embargo, el artefacto operativo aún debía "
    "materializarse en el proyecto, ya que el TSV manual disponible contenía solo encabezados y 0 filas.\n\n"
    "Este proceso materializa un artefacto operativo controlado en Excel y CSV, sin modificar la hoja original, "
    "sin actualizar el TSV base y sin automatizar sustituciones. La base final queda en el grano CLAVE_OPERATIVA, "
    "con 59 registros de Pregrado, 0 códigos SIES vacíos, 0 claves duplicadas y validación fuerte sobre Código SIES "
    "y sus derivados.\n\n"
    "La advertencia REVISAR_COD_CARRERA no bloquea el uso de la base, porque Cód. Carrera CNED y COD_CAR_DERIVADO "
    "pertenecen a naturalezas distintas. La validación operativa debe realizarse sobre Código SIES, COD_SED_DERIVADO, "
    "COD_CAR_DERIVADO, JOR_DERIVADA, VERSION_DERIVADA y CLAVE_OPERATIVA."
)


@dataclass(frozen=True)
class Paths:
    repo: Path
    cned: Path
    data_dir: Path
    resultados_dir: Path
    script: Path
    script_out: Path
    excel_out: Path
    csv_out: Path
    md_out: Path


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def build_paths(script: Path, timestamp: str) -> Paths:
    repo = script.parents[3]
    cned = repo / "indices_2025" / "cned"
    data_dir = cned / "data"
    resultados_dir = cned / "resultados"
    stem = f"CNED_ARTEFACTO_OPERATIVO_LIMPIO_{timestamp}"
    return Paths(
        repo=repo,
        cned=cned,
        data_dir=data_dir,
        resultados_dir=resultados_dir,
        script=script,
        script_out=script,
        excel_out=resultados_dir / f"{stem}.xlsx",
        csv_out=data_dir / "BASE_CNED_LIMPIA_OPERATIVA.csv",
        md_out=resultados_dir / f"{stem}.md",
    )


def ensure_dirs(paths: Paths) -> None:
    paths.data_dir.mkdir(parents=True, exist_ok=True)
    paths.resultados_dir.mkdir(parents=True, exist_ok=True)


def load_embedded_tsv(text: str) -> pd.DataFrame:
    df = pd.read_csv(StringIO(text), sep="\t", dtype=str, keep_default_na=False)
    df = df.fillna("")
    df.columns = [clean(col) for col in df.columns]
    for col in df.columns:
        df[col] = df[col].map(clean)
    return df


def build_resumen_df(timestamp: str) -> pd.DataFrame:
    rows = [
        ("Fuente", "Hoja1", "", ""),
        ("Generado", timestamp, "", ""),
        ("", "", "", ""),
        ("CONTEOS GLOBALES", "", "", ""),
        ("Métrica", "Valor", "Detalle", ""),
        ("Total de filas originales", "354", "Excluye encabezado", ""),
        ("Total de filas con Código SIES vacío", "7", "Enviadas a SIN_CODIGO_SIES", ""),
        ("Total de filas con Código SIES válido", "347", "No vacío; formato se valida aparte", ""),
        ("Total Código SIES con formato OK", "347", "Formato I162S#C#J#V#", ""),
        ("Total Código SIES con formato REVISAR", "0", "No cumple formato esperado", ""),
        ("Total de claves operativas únicas", "59", "Código SIES | Cód. Carrera | Nombre Programa | Horario", ""),
        ("Total de duplicados eliminados", "288", "Registrados en DUPLICADOS_CNED", ""),
        ("Total de registros finales en BASE_CNED_LIMPIA", "59", "Un registro por CLAVE_OPERATIVA", ""),
        (
            "Años presentes en la base original",
            "2005, 2006, 2007, 2008, 2009, 2010, 2011, 2012, 2013, 2014, 2015, 2016, 2017, 2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025",
            "",
            "",
        ),
        ("Año mínimo", "2005", "Base original", ""),
        ("Año máximo", "2025", "Base original", ""),
        ("", "", "", ""),
        ("CANTIDAD DE REGISTROS POR AÑO - BASE ORIGINAL", "", "", ""),
        ("Año", "Registros", "", ""),
        ("2005", "8", "", ""),
        ("2006", "8", "", ""),
        ("2007", "8", "", ""),
        ("2008", "8", "", ""),
        ("2009", "8", "", ""),
        ("2010", "8", "", ""),
        ("2011", "8", "", ""),
        ("2012", "8", "", ""),
        ("2013", "12", "", ""),
        ("2014", "13", "", ""),
        ("2015", "13", "", ""),
        ("2016", "16", "", ""),
        ("2017", "18", "", ""),
        ("2018", "22", "", ""),
        ("2019", "28", "", ""),
        ("2020", "32", "", ""),
        ("2021", "31", "", ""),
        ("2022", "28", "", ""),
        ("2023", "18", "", ""),
        ("2024", "25", "", ""),
        ("2025", "34", "", ""),
        ("", "", "", ""),
        ("CANTIDAD DE REGISTROS POR PREGRADO/POSGRADO - BASE FINAL", "", "", ""),
        ("Pregrado/Posgrado", "Registros", "", ""),
        ("Pregrado", "59", "", ""),
        ("", "", "", ""),
        ("CANTIDAD DE REGISTROS POR HORARIO - BASE FINAL", "", "", ""),
        ("Horario", "Registros", "", ""),
        ("Diurno", "14", "", ""),
        ("Otro", "22", "", ""),
        ("Vespertino", "23", "", ""),
        ("", "", "", ""),
        ("CANTIDAD DE REGISTROS POR JOR_DERIVADA - BASE FINAL", "", "", ""),
        ("JOR_DERIVADA", "Registros", "", ""),
        ("1", "14", "", ""),
        ("2", "23", "", ""),
        ("4", "22", "", ""),
        ("", "", "", ""),
        ("VALIDACIONES EN BASE FINAL", "", "", ""),
        ("Validación", "Resultado", "Registros", ""),
        ("VALIDACION_CODIGO_UNICO", "OK", "59", ""),
        ("VALIDACION_COD_CARRERA", "REVISAR_COD_CARRERA", "59", ""),
    ]
    return pd.DataFrame(rows, columns=["RESUMEN_LIMPIEZA", "Columna1", "Columna2", "Columna3"])


def build_duplicados_df() -> pd.DataFrame:
    rows = [
        {
            "MENSAJE": "Detalle completo de DUPLICADOS_CNED no fue pegado en este prompt",
            "TOTAL_DUPLICADOS_DECLARADOS": str(EXPECTED_DUPLICADOS),
            "CRITERIO_DUPLICIDAD": "Duplicidad operativa por CLAVE_OPERATIVA",
            "OBSERVACION": "Hoja creada para trazabilidad; completar con detalle si se requiere auditoría fila a fila",
        }
    ]
    return pd.DataFrame(rows)


def dict_counts(series: pd.Series) -> dict[str, int]:
    counts = series.map(clean).value_counts(dropna=False).to_dict()
    return {str(key): int(value) for key, value in counts.items() if clean(key)}


def validate_frames(base_df: pd.DataFrame, sin_df: pd.DataFrame) -> tuple[dict[str, Any], list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []

    if base_df.shape[0] != EXPECTED_BASE_ROWS:
        errors.append(f"BASE_CNED_LIMPIA debe tener {EXPECTED_BASE_ROWS} filas y tiene {base_df.shape[0]}")
    if sin_df.shape[0] != EXPECTED_SIN_ROWS:
        errors.append(f"SIN_CODIGO_SIES debe tener {EXPECTED_SIN_ROWS} filas y tiene {sin_df.shape[0]}")
    if base_df.shape[1] != EXPECTED_BASE_COLS:
        errors.append(f"BASE_CNED_LIMPIA debe tener {EXPECTED_BASE_COLS} columnas y tiene {base_df.shape[1]}")
    if sin_df.shape[1] != EXPECTED_SIN_COLS:
        errors.append(f"SIN_CODIGO_SIES debe tener {EXPECTED_SIN_COLS} columnas y tiene {sin_df.shape[1]}")

    required_base = [
        "Código SIES",
        "CLAVE_OPERATIVA",
        "VALIDACION_CODIGO_UNICO",
        "VALIDACION_COD_CARRERA",
        "Pregrado/Posgrado",
        "Horario",
        "JOR_DERIVADA",
    ]
    required_sin = ["Código SIES", "Pregrado/Posgrado", "MOTIVO_REVISION"]

    missing_base = [col for col in required_base if col not in base_df.columns]
    missing_sin = [col for col in required_sin if col not in sin_df.columns]
    if missing_base:
        errors.append("Faltan columnas en BASE_CNED_LIMPIA: " + ", ".join(missing_base))
    if missing_sin:
        errors.append("Faltan columnas en SIN_CODIGO_SIES: " + ", ".join(missing_sin))

    base_codigo_vacio = int(base_df.get("Código SIES", pd.Series(dtype=str)).map(clean).eq("").sum())
    sin_codigo_vacio = int(sin_df.get("Código SIES", pd.Series(dtype=str)).map(clean).eq("").sum())
    clave_unicas = int(base_df.get("CLAVE_OPERATIVA", pd.Series(dtype=str)).map(clean).nunique())
    validacion_codigo = dict_counts(base_df.get("VALIDACION_CODIGO_UNICO", pd.Series(dtype=str)))
    validacion_carrera = dict_counts(base_df.get("VALIDACION_COD_CARRERA", pd.Series(dtype=str)))
    pregrado_counts = dict_counts(base_df.get("Pregrado/Posgrado", pd.Series(dtype=str)))
    horario_counts = dict_counts(base_df.get("Horario", pd.Series(dtype=str)))
    jor_counts = dict_counts(base_df.get("JOR_DERIVADA", pd.Series(dtype=str)).astype(str))
    sin_pregrado_counts = dict_counts(sin_df.get("Pregrado/Posgrado", pd.Series(dtype=str)))
    sin_motivo_counts = dict_counts(sin_df.get("MOTIVO_REVISION", pd.Series(dtype=str)))

    codigo_patron = re.compile(r"I162S\d+C\d+J\d+V\d+")
    codigo_ok = int(base_df.get("Código SIES", pd.Series(dtype=str)).map(clean).map(lambda value: bool(codigo_patron.fullmatch(value))).sum())

    if base_codigo_vacio != 0:
        errors.append(f"BASE_CNED_LIMPIA debe tener 0 Código SIES vacíos y tiene {base_codigo_vacio}")
    if sin_codigo_vacio != EXPECTED_SIN_ROWS:
        errors.append(f"SIN_CODIGO_SIES debe tener {EXPECTED_SIN_ROWS} Código SIES vacíos y tiene {sin_codigo_vacio}")
    if clave_unicas != EXPECTED_BASE_ROWS:
        errors.append(f"BASE_CNED_LIMPIA debe tener {EXPECTED_BASE_ROWS} CLAVE_OPERATIVA únicas y tiene {clave_unicas}")
    if validacion_codigo != EXPECTED_VALIDACION_CODIGO:
        errors.append(f"VALIDACION_CODIGO_UNICO esperado {EXPECTED_VALIDACION_CODIGO} y obtenido {validacion_codigo}")
    if validacion_carrera != EXPECTED_VALIDACION_CARRERA:
        errors.append(f"VALIDACION_COD_CARRERA esperado {EXPECTED_VALIDACION_CARRERA} y obtenido {validacion_carrera}")
    if pregrado_counts != EXPECTED_PREGRADO:
        errors.append(f"Pregrado/Posgrado esperado {EXPECTED_PREGRADO} y obtenido {pregrado_counts}")
    if horario_counts != EXPECTED_HORARIO:
        errors.append(f"Horario esperado {EXPECTED_HORARIO} y obtenido {horario_counts}")
    if jor_counts != EXPECTED_JOR:
        errors.append(f"JOR_DERIVADA esperado {EXPECTED_JOR} y obtenido {jor_counts}")
    if codigo_ok != EXPECTED_BASE_ROWS:
        errors.append(f"Código SIES con patrón válido esperado {EXPECTED_BASE_ROWS} y obtenido {codigo_ok}")
    if sin_pregrado_counts != EXPECTED_SIN_PREGRADO:
        errors.append(f"Pregrado/Posgrado en SIN_CODIGO_SIES esperado {EXPECTED_SIN_PREGRADO} y obtenido {sin_pregrado_counts}")
    if sin_motivo_counts != EXPECTED_SIN_MOTIVO:
        errors.append(f"MOTIVO_REVISION esperado {EXPECTED_SIN_MOTIVO} y obtenido {sin_motivo_counts}")

    warnings.append("REVISAR_COD_CARRERA se mantiene como advertencia no bloqueante en los 59 registros de la base final")
    warnings.append("DUPLICADOS_CNED se materializa como hoja de trazabilidad sin detalle fila a fila; el resumen declara 288 duplicados eliminados")

    metrics = {
        "base_rows": int(base_df.shape[0]),
        "sin_rows": int(sin_df.shape[0]),
        "base_columns": int(base_df.shape[1]),
        "sin_columns": int(sin_df.shape[1]),
        "base_codigo_vacio": base_codigo_vacio,
        "sin_codigo_vacio": sin_codigo_vacio,
        "claves_operativas_unicas": clave_unicas,
        "codigo_sies_formato_ok": codigo_ok,
        "validacion_codigo": validacion_codigo,
        "validacion_carrera": validacion_carrera,
        "pregrado_counts": pregrado_counts,
        "horario_counts": horario_counts,
        "jor_counts": jor_counts,
        "sin_pregrado_counts": sin_pregrado_counts,
        "sin_motivo_counts": sin_motivo_counts,
        "duplicados_declarados": EXPECTED_DUPLICADOS,
    }
    return metrics, warnings, errors


def adjust_widths(ws: Any) -> None:
    for col_idx in range(1, ws.max_column + 1):
        values = [clean(ws.cell(row=row_idx, column=col_idx).value) for row_idx in range(1, min(ws.max_row, 300) + 1)]
        width = min(max(max((len(value) for value in values), default=0) + 2, 12), 60)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def apply_base_highlights(ws: Any) -> None:
    header_map = {clean(ws.cell(1, col_idx).value): col_idx for col_idx in range(1, ws.max_column + 1)}
    cod_col = header_map.get("VALIDACION_CODIGO_UNICO")
    carr_col = header_map.get("VALIDACION_COD_CARRERA")
    if cod_col is not None:
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, cod_col).fill = copy(OK_FILL)
    if carr_col is not None:
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, carr_col).fill = copy(WARN_FILL)


def apply_summary_sections(ws: Any) -> None:
    for row_idx in range(2, ws.max_row + 1):
        first = clean(ws.cell(row_idx, 1).value)
        second = clean(ws.cell(row_idx, 2).value)
        third = clean(ws.cell(row_idx, 3).value)
        fourth = clean(ws.cell(row_idx, 4).value)
        if first and not second and not third and not fourth:
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row_idx, col_idx)
                cell.fill = copy(SECTION_FILL)
                cell.font = Font(bold=True)


def write_dataframe_sheet(wb: Workbook, title: str, df: pd.DataFrame, kind: str) -> None:
    ws = wb.create_sheet(title)
    for row in dataframe_to_rows(df.fillna(""), index=False, header=True):
        ws.append(row)

    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(1, col_idx)
        cell.fill = copy(HEADER_FILL)
        cell.font = copy(HEADER_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if kind == "base":
        apply_base_highlights(ws)
    elif kind == "summary":
        apply_summary_sections(ws)
    elif kind == "duplicates" and ws.max_row >= 2:
        for col_idx in range(1, ws.max_column + 1):
            ws.cell(2, col_idx).fill = copy(NOTE_FILL)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 1)}"
    adjust_widths(ws)


def build_markdown(paths: Paths, timestamp: str, metrics: dict[str, Any], warnings: list[str]) -> str:
    horario = metrics["horario_counts"]
    jor = metrics["jor_counts"]
    lines = [
        "# CNED Artefacto Operativo Limpio",
        "",
        "## Dictamen",
        "",
        DICTAMEN,
        "",
        "## Archivos generados",
        "",
        f"- Excel: {paths.excel_out}",
        f"- CSV: {paths.csv_out}",
        f"- Markdown: {paths.md_out}",
        f"- Script: {paths.script_out}",
        "",
        "## Conteos",
        "",
        f"- BASE_CNED_LIMPIA: {metrics['base_rows']} registros",
        f"- SIN_CODIGO_SIES: {metrics['sin_rows']} registros",
        f"- Código SIES vacíos en BASE_CNED_LIMPIA: {metrics['base_codigo_vacio']}",
        f"- CLAVE_OPERATIVA únicas: {metrics['claves_operativas_unicas']}",
        f"- VALIDACION_CODIGO_UNICO: {metrics['validacion_codigo']}",
        f"- VALIDACION_COD_CARRERA: {metrics['validacion_carrera']}",
        f"- Horario: {horario}",
        f"- JOR_DERIVADA: {jor}",
        f"- DUPLICADOS declarados: {metrics['duplicados_declarados']}",
        "",
        "## Advertencias",
        "",
    ]
    lines.extend(f"- {warning}" for warning in warnings)
    lines.extend(
        [
            "- No se modificó Hoja1 ni se actualizó DATOS_VALIDACION_CNED_TSV.tsv.",
            "- No se actualizó DATOS_VALIDACION_CNED_TSV_ACTUALIZADO.tsv.",
            "- BASE_CNED_LIMPIA_OPERATIVA.csv se exportó como respaldo controlado de la hoja principal.",
            "",
            "## Reglas de uso",
            "",
            "- Usar BASE_CNED_LIMPIA solo en el grano CLAVE_OPERATIVA.",
            "- Validar operativamente con Código SIES, COD_SED_DERIVADO, COD_CAR_DERIVADO, JOR_DERIVADA y VERSION_DERIVADA.",
            "- No interpretar REVISAR_COD_CARRERA como error bloqueante.",
            "- No mezclar SIN_CODIGO_SIES con la base operativa principal.",
            "- No usar este artefacto para reemplazar automáticamente TSV base ni automatizar sustituciones.",
            "",
            "## Cierre",
            "",
            CIERRE_TEXTO,
            "",
            f"Generado: {timestamp}",
        ]
    )
    return "\n".join(lines) + "\n"


def main() -> int:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    script = Path(__file__).resolve()
    paths = build_paths(script, timestamp)
    ensure_dirs(paths)

    base_df = load_embedded_tsv(BASE_CNED_LIMPIA_TSV)
    sin_df = load_embedded_tsv(SIN_CODIGO_SIES_TSV)
    duplicados_df = build_duplicados_df()
    resumen_df = build_resumen_df(timestamp)

    metrics, warnings, errors = validate_frames(base_df, sin_df)
    if errors:
        print("ERROR: validación fallida del artefacto operativo CNED limpio")
        for item in errors:
            print(f"- {item}")
        return 1

    workbook = Workbook()
    workbook.remove(workbook.active)
    write_dataframe_sheet(workbook, "RESUMEN_LIMPIEZA", resumen_df, "summary")
    write_dataframe_sheet(workbook, "BASE_CNED_LIMPIA", base_df, "base")
    write_dataframe_sheet(workbook, "SIN_CODIGO_SIES", sin_df, "sin")
    write_dataframe_sheet(workbook, "DUPLICADOS_CNED", duplicados_df, "duplicates")
    workbook.save(paths.excel_out)

    base_df.to_csv(paths.csv_out, sep=";", index=False, encoding="utf-8")
    paths.md_out.write_text(build_markdown(paths, timestamp, metrics, warnings), encoding="utf-8")

    print("=" * 120)
    print("CNED — Materialización de artefacto operativo limpio")
    print("=" * 120)
    print(f"Excel generado: {paths.excel_out}")
    print(f"CSV generado:   {paths.csv_out}")
    print(f"Markdown:       {paths.md_out}")
    print(f"Script:         {paths.script_out}")
    print()
    print(f"DICTAMEN FINAL: {DICTAMEN}")
    print()
    print("Validaciones principales:")
    print(f"- BASE_CNED_LIMPIA filas: {metrics['base_rows']}")
    print(f"- SIN_CODIGO_SIES filas: {metrics['sin_rows']}")
    print(f"- Código SIES vacíos en BASE_CNED_LIMPIA: {metrics['base_codigo_vacio']}")
    print(f"- CLAVE_OPERATIVA únicas: {metrics['claves_operativas_unicas']}")
    print(f"- VALIDACION_CODIGO_UNICO: {metrics['validacion_codigo']}")
    print(f"- VALIDACION_COD_CARRERA: {metrics['validacion_carrera']} (advertencia no bloqueante)")
    print(f"- Pregrado/Posgrado BASE: {metrics['pregrado_counts']}")
    print(f"- Horario BASE: {metrics['horario_counts']}")
    print(f"- JOR_DERIVADA BASE: {metrics['jor_counts']}")
    print(f"- DUPLICADOS declarados: {metrics['duplicados_declarados']}")
    print()
    print("Advertencias:")
    for warning in warnings:
        print(f"- {warning}")
    print("=" * 120)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''