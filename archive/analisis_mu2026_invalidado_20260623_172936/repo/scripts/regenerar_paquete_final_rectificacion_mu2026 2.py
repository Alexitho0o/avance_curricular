#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import importlib.util
import shutil
import subprocess
import sys
import textwrap
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


BASE = Path("/Users/alexi/Documents/GitHub/avance_curricular")
OUT_DIR = BASE / "resultados/auditoria_multicodcli_reconstruida_2026/cierre_34_multivigentes_sies_20260623"
DESKTOP = Path("/Users/alexi/Desktop")
SCRIPTS_DIR = BASE / "scripts"

AUDITABLE_SCRIPT = SCRIPTS_DIR / "generar_cierre_34_multivigentes_sies_auditable.py"
RECTIFICATION_HELPER_SCRIPT = SCRIPTS_DIR / "generar_rectificacion_matricula_unificada_pregrado_2026.py"
MASTER_SCRIPT = SCRIPTS_DIR / "regenerar_paquete_final_rectificacion_mu2026.py"

AUDITABLE_XLSX = OUT_DIR / "CIERRE_34_MULTIVIGENTES_SIES_AUDITABLE.xlsx"
TECHNICAL_XLSX = OUT_DIR / "CIERRE_34_MULTIVIGENTES_SIES.xlsx"
SUMMARY_XLSX = OUT_DIR / "RESUMEN_EJECUTIVO_CIERRE_34_MULTIVIGENTES_SIES.xlsx"
RECTIFICATION_CSV = OUT_DIR / "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026.csv"
RECTIFICATION_XLSX = OUT_DIR / "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026_CON_ENCABEZADOS.xlsx"
EMAIL_BODY_TXT = OUT_DIR / "CUERPO_CORREO_RECTORIA_RECTIFICACION_MU2026.txt"

DESKTOP_AUDITABLE_XLSX = DESKTOP / AUDITABLE_XLSX.name
DESKTOP_SUMMARY_XLSX = DESKTOP / SUMMARY_XLSX.name
DESKTOP_RECTIFICATION_CSV = DESKTOP / RECTIFICATION_CSV.name
DESKTOP_RECTIFICATION_XLSX = DESKTOP / RECTIFICATION_XLSX.name
DESKTOP_EMAIL_BODY_TXT = DESKTOP / EMAIL_BODY_TXT.name

PROMEDIOS_XLSX = BASE / "input/PROMEDIOSDEALUMNOS_7804.xlsx"
MANUAL_PDF = BASE / "Manual_Matrícula_Unificada_2026.pdf"
OFFICIAL_CSV = BASE / "control/punto_0_carga_principal_mu2026/archivos_congelados/PARA_SUBIR_DESKTOP__matricula_unificada_2026_pregrado_PARA_SUBIR.csv"
PES_READY_CSV = BASE / "resultados/matricula_unificada_2026_pregrado_PES_READY.csv"
HEADER_CSV = BASE / "resultados/matricula_unificada_2026_control.csv"
CESAR_SOURCE_CSV = BASE / "resultados/matricula_unificada_rectificacion_20260618/matricula_unificada_2026_BASE_COMPLETA_CORREGIDA_P1_20260618.csv"
REVISION_34_XLSX = BASE / "resultados/auditoria_multicodcli_reconstruida_2026/cierre_brechas_20260622/REVISION_34_RUT_MULTIVIGENTES.xlsx"
TRAZABILIDAD_CIERRE_XLSX = BASE / "resultados/auditoria_multicodcli_reconstruida_2026/cierre_brechas_20260622/TRAZABILIDAD_CIERRE_RUT_CODCLI.xlsx"

MU32_COLUMNS = [
    "TIPO_DOC", "N_DOC", "DV", "PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRE",
    "SEXO", "FECH_NAC", "NAC", "PAIS_EST_SEC", "COD_SED", "COD_CAR", "MODALIDAD",
    "JOR", "VERSION", "FOR_ING_ACT", "ANIO_ING_ACT", "SEM_ING_ACT", "ANIO_ING_ORI",
    "SEM_ING_ORI", "ASI_INS_ANT", "ASI_APR_ANT", "PROM_PRI_SEM", "PROM_SEG_SEM",
    "ASI_INS_HIS", "ASI_APR_HIS", "NIV_ACA", "SIT_FON_SOL", "SUS_PRE",
    "FECHA_MATRICULA", "REINCORPORACION", "VIG",
]

FINAL_CASES = [
    {
        "rut": "17726298",
        "student": "José Jairo Velásquez Figueroa",
        "student_upper": "JOSE JAIRO VELASQUEZ FIGUEROA",
        "action": "INCORPORAR",
        "codcli_or_record": "20253IINF012",
        "codcli": "20253IINF012",
        "offer": "S2C1M3J4V2",
        "vig": "1",
    },
    {
        "rut": "18059242",
        "student": "César Andrés Rubilar Sanhueza",
        "student_upper": "CESAR ANDRES RUBILAR SANHUEZA",
        "action": "ELIMINAR/ANULAR",
        "codcli_or_record": "MU_LINEA_04073|RUT=18059242-3|S2C1M1J1V1",
        "codcli": "",
        "offer": "S2C1M1J1V1",
        "vig": "0",
    },
    {
        "rut": "18939583",
        "student": "Adolfo Andrés Campos Gómez",
        "student_upper": "ADOLFO ANDRES CAMPOS GOMEZ",
        "action": "INCORPORAR",
        "codcli_or_record": "20241TPAS250",
        "codcli": "20241TPAS250",
        "offer": "S2C57M3J4V1",
        "vig": "1",
    },
    {
        "rut": "19356713",
        "student": "Nicolás Octavio Lagos Vega",
        "student_upper": "NICOLAS OCTAVIO LAGOS VEGA",
        "action": "INCORPORAR",
        "codcli_or_record": "20251ICDA004",
        "codcli": "20251ICDA004",
        "offer": "S2C88M3J4V1",
        "vig": "1",
    },
    {
        "rut": "19513133",
        "student": "Ariel Eduardo Martínez Álvarez",
        "student_upper": "ARIEL EDUARDO MARTINEZ ALVAREZ",
        "action": "INCORPORAR",
        "codcli_or_record": "20261ICDA007",
        "codcli": "20261ICDA007",
        "offer": "S2C88M3J4V1",
        "vig": "1",
    },
]

MARCOS_RUT = "15651488"
MARCOS_CODCLI = "20261DDASC008"
FORBIDDEN_EXECUTIVE_TOKENS = [
    MARCOS_RUT,
    MARCOS_CODCLI,
    "MARCOS ANTONIO QUEZADA MILLAHUAL",
    "QUEZADA MILLAHUAL",
    "NO_DETERMINABLE",
    "NO_INFORMAR",
]

AUDITABLE_EXPECTED_FORMULAS = {
    "00_RESUMEN": 22,
    "01_CASOS_TRAZABLES": 476,
    "02_DETALLE_CODCLI": 952,
    "04_FILAS_MU": 36,
    "09_CONTROLES": 18,
}


def fail(message: str) -> None:
    raise RuntimeError(message)


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat", "<na>"}:
        return ""
    if text.endswith(".0") and text[:-2].lstrip("-").isdigit():
        return text[:-2]
    return " ".join(text.split())


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def backup_if_exists(path: Path, backups: list[Path]) -> None:
    if not path.exists():
        return
    backup = path.with_name(f"{path.name}.bak_{timestamp()}")
    shutil.copy2(path, backup)
    backups.append(backup)


def run_checked(cmd: list[str], *, cwd: Path = BASE, label: str) -> subprocess.CompletedProcess[str]:
    print(f"\nEJECUTANDO: {label}")
    print(" ".join(cmd))
    result = subprocess.run(cmd, cwd=cwd, text=True, capture_output=True)
    if result.stdout:
        print(result.stdout)
    if result.stderr:
        print(result.stderr, file=sys.stderr)
    if result.returncode != 0:
        fail(f"{label} terminó con exit {result.returncode}")
    return result


def import_rectification_helper() -> Any:
    spec = importlib.util.spec_from_file_location("rectification_helper", RECTIFICATION_HELPER_SCRIPT)
    if spec is None or spec.loader is None:
        fail(f"No se pudo cargar helper de rectificación: {RECTIFICATION_HELPER_SCRIPT}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def source_paths() -> list[Path]:
    return [
        TECHNICAL_XLSX,
        PROMEDIOS_XLSX,
        MANUAL_PDF,
        OFFICIAL_CSV,
        PES_READY_CSV,
        HEADER_CSV,
        CESAR_SOURCE_CSV,
        REVISION_34_XLSX,
        TRAZABILIDAD_CIERRE_XLSX,
        AUDITABLE_SCRIPT,
        RECTIFICATION_HELPER_SCRIPT,
    ]


def snapshot_hashes(paths: list[Path]) -> dict[str, str]:
    out = {}
    for path in paths:
        if not path.exists():
            fail(f"No existe fuente requerida: {path}")
        out[str(path)] = sha256(path)
    return out


def recalc_with_microsoft_excel(path: Path) -> None:
    applescript = textwrap.dedent(
        f"""
        set targetFile to "{path}"
        tell application "Microsoft Excel"
            set display alerts to false
            try
                close workbook "{path.name}" saving yes
            end try
            set wb to open workbook workbook file name targetFile
            calculate full rebuild
            delay 8
            save wb
            close workbook "{path.name}" saving yes
        end tell
        """
    ).strip()
    run_checked(["osascript", "-e", applescript], label=f"recalcular con Microsoft Excel {path.name}")


def formula_counts(path: Path) -> dict[str, int]:
    wb = load_workbook(path, read_only=True, data_only=False)
    counts: dict[str, int] = {}
    for ws in wb.worksheets:
        total = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    total += 1
        counts[ws.title] = total
    wb.close()
    return counts


def rows_from_ws(ws: Any) -> list[dict[str, Any]]:
    headers = [clean(cell.value) for cell in ws[1]]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(clean(v) for v in row):
            continue
        rows.append({headers[idx]: row[idx] if idx < len(row) else "" for idx in range(len(headers))})
    return rows


def validate_no_external_links(path: Path) -> tuple[int, int]:
    with zipfile.ZipFile(path) as zf:
        external_links = [n for n in zf.namelist() if n.startswith("xl/externalLinks/")]
        external_formulas = []
        for name in zf.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                data = zf.read(name).decode("utf-8", errors="ignore")
                if "<f>" in data and "[" in data and "]" in data:
                    external_formulas.append(name)
        return len(external_links), len(external_formulas)


def validate_auditable(rect_helper: Any) -> dict[str, Any]:
    if not AUDITABLE_XLSX.exists():
        fail(f"No existe auditable regenerado: {AUDITABLE_XLSX}")
    wb_formula = load_workbook(AUDITABLE_XLSX, read_only=True, data_only=False)
    sheetnames = wb_formula.sheetnames
    expected_sheets = [
        "00_RESUMEN",
        "01_CASOS_TRAZABLES",
        "02_DETALLE_CODCLI",
        "03_EQUIVALENCIAS",
        "04_FILAS_MU",
        "05_RAW_DATOSALUMNOS",
        "06_RAW_MATRIZ",
        "07_RAW_MU",
        "08_REGLAS",
        "09_CONTROLES",
    ]
    if sheetnames != expected_sheets:
        fail(f"Auditable no tiene las 10 hojas esperadas: {sheetnames}")
    counts = formula_counts(AUDITABLE_XLSX)
    if counts != {**{s: 0 for s in expected_sheets}, **AUDITABLE_EXPECTED_FORMULAS}:
        expected = {**{s: 0 for s in expected_sheets}, **AUDITABLE_EXPECTED_FORMULAS}
        fail(f"Conteo de fórmulas auditable inválido: {counts}; esperado {expected}")
    total_formulas = sum(counts.values())
    if total_formulas != 1504:
        fail(f"Auditable tiene {total_formulas} fórmulas, esperado 1504")
    wb_formula.close()

    wb = load_workbook(AUDITABLE_XLSX, read_only=True, data_only=True)
    raw_da = wb["05_RAW_DATOSALUMNOS"]
    raw_matriz = wb["06_RAW_MATRIZ"]
    if (raw_da.max_row, raw_da.max_column) != (13708, 74):
        fail(f"05_RAW_DATOSALUMNOS dimensión inválida: {raw_da.max_row}x{raw_da.max_column}")
    if (raw_matriz.max_row, raw_matriz.max_column) != (140, 19):
        fail(f"06_RAW_MATRIZ dimensión inválida: {raw_matriz.max_row}x{raw_matriz.max_column}")

    cases = rows_from_ws(wb["01_CASOS_TRAZABLES"])
    detail = rows_from_ws(wb["02_DETALLE_CODCLI"])
    controls = rows_from_ws(wb["09_CONTROLES"])
    if len(cases) != 34:
        fail(f"Auditable casos esperados 34, obtenido {len(cases)}")
    if len(detail) != 68:
        fail(f"Auditable detalle CODCLI esperado 68, obtenido {len(detail)}")

    class_counts = Counter(clean(r.get("CLASIFICACION_FINAL")) for r in cases)
    expected_class = {"CORRECCION_SIES": 5, "NO_INFORMAR": 28, "NO_DETERMINABLE": 1}
    for key, expected in expected_class.items():
        if class_counts.get(key, 0) != expected:
            fail(f"Auditable clasificación {key}: {class_counts.get(key, 0)} != {expected}")

    es_2026 = Counter(clean(r.get("ES_2026_1")) for r in detail)
    elegible = sum(1 for r in detail if clean(r.get("ELEGIBLE_FINAL")) in {"1", "SI"})
    mapeo = Counter(clean(r.get("MAPEO_DEMOSTRADO")) for r in detail)
    evidence = Counter(clean(r.get("EVIDENCIA_COMPLETA")) for r in detail)
    if es_2026.get("SI", 0) != 68:
        fail(f"Auditable ES_2026_1 no devuelve SI en 68 filas: {dict(es_2026)}")
    if elegible != 68:
        fail(f"Auditable CODCLI elegibles esperado 68, obtenido {elegible}")
    if mapeo.get("SI", 0) != 67 or mapeo.get("NO", 0) != 1:
        fail(f"Auditable MAPEO_DEMOSTRADO inválido: {dict(mapeo)}")
    marcos = [r for r in detail if clean(r.get("CODCLI")) == MARCOS_CODCLI]
    if len(marcos) != 1 or clean(marcos[0].get("MAPEO_DEMOSTRADO")) != "NO":
        fail("Auditable no conserva a Marcos Quezada como único mapeo SIES no demostrado")
    if evidence.get("SI", 0) != 68:
        fail(f"Auditable evidencia completa inválida: {dict(evidence)}")

    control_bad = [r for r in controls if clean(r.get("ESTADO")) and clean(r.get("ESTADO")) != "OK"]
    if control_bad:
        fail(f"Auditable tiene controles no OK: {control_bad[:5]}")

    correction_without_mapping = sum(
        1 for r in cases
        if clean(r.get("CLASIFICACION_FINAL")) == "CORRECCION_SIES" and clean(r.get("MAPEO_COMPLETO")) == "NO"
    )
    correction_without_evidence = sum(
        1 for r in cases
        if clean(r.get("CLASIFICACION_FINAL")) == "CORRECCION_SIES" and clean(r.get("EVIDENCIA_COMPLETA")) == "NO"
    )
    duplicated_rut = len(cases) - len({clean(r.get("RUT_NORM")) for r in cases})
    if correction_without_mapping or correction_without_evidence or duplicated_rut:
        fail(
            "Auditable controles básicos fallan: "
            f"duplicados={duplicated_rut}, sin_mapeo={correction_without_mapping}, sin_evidencia={correction_without_evidence}"
        )
    wb.close()

    casos_df, detalle_df, _ = rect_helper.load_auditable()
    no_informar_ruts, no_det_ruts = rect_helper.validate_decision_sources(casos_df, detalle_df)
    if no_det_ruts != {MARCOS_RUT}:
        fail(f"NO_DETERMINABLE inesperados: {sorted(no_det_ruts)}")
    tech_cases = pd.read_excel(TECHNICAL_XLSX, sheet_name="CASOS", dtype=str).fillna("")
    tech_cases.columns = [clean(c) for c in tech_cases.columns]
    tech_cases["RUT"] = tech_cases["RUT"].map(rect_helper.rut_norm)
    tech_cases["CLASIFICACION_TECNICA_TRADUCIDA"] = tech_cases.apply(
        lambda r: "CORRECCION_SIES"
        if clean(r.get("REQUIERE_INFORMAR_SIES")) == "SI"
        else ("NO_DETERMINABLE" if clean(r.get("DECISION_FINAL")).startswith("NO_DETERMINABLE") else "NO_INFORMAR"),
        axis=1,
    )
    audit_map = {rect_helper.rut_norm(r["RUT_NORM"]): clean(r["CLASIFICACION_FINAL"]) for _, r in casos_df.iterrows()}
    tech_map = dict(zip(tech_cases["RUT"], tech_cases["CLASIFICACION_TECNICA_TRADUCIDA"]))
    differences = [
        {"RUT": rut, "AUDITABLE": audit_map.get(rut), "TECNICO": tech_map.get(rut)}
        for rut in sorted(set(audit_map) | set(tech_map))
        if audit_map.get(rut) != tech_map.get(rut)
    ]
    if differences:
        fail(f"Diferencias reales contra cierre técnico: {differences[:10]}")

    external_links, external_formulas = validate_no_external_links(AUDITABLE_XLSX)
    if external_links or external_formulas:
        fail(f"Auditable tiene vínculos externos: links={external_links}, formulas={external_formulas}")

    return {
        "sheetnames": sheetnames,
        "formula_counts": counts,
        "total_formulas": total_formulas,
        "rut_total": len(cases),
        "detail_total": len(detail),
        "class_counts": dict(class_counts),
        "es_2026": dict(es_2026),
        "eligible_codcli": elegible,
        "no_informar_ruts": no_informar_ruts,
        "no_determinable_ruts": no_det_ruts,
        "differences_against_technical": differences,
        "external_links": external_links,
        "external_formulas": external_formulas,
    }


def build_rectification_records(rect_helper: Any, template: Any, headers: list[str]) -> tuple[pd.DataFrame, pd.DataFrame, list[dict[str, str]]]:
    casos_df, detalle_df, _ = rect_helper.load_auditable()
    no_informar_ruts, no_det_ruts = rect_helper.validate_decision_sources(casos_df, detalle_df)
    if no_det_ruts != {MARCOS_RUT}:
        fail(f"NO_DETERMINABLE inesperados al construir rectificación: {sorted(no_det_ruts)}")

    datos = pd.read_excel(PROMEDIOS_XLSX, sheet_name="DatosAlumnos", dtype=str)
    h1 = pd.read_excel(PROMEDIOS_XLSX, sheet_name="Hoja1", dtype=str)
    datos.columns = [clean(c) for c in datos.columns]
    h1.columns = [clean(c) for c in h1.columns]

    records: list[dict[str, str]] = []
    trace_rows: list[dict[str, str]] = []
    cesar_diff: list[dict[str, str]] = []

    for idx, case in enumerate(FINAL_CASES, start=1):
        if case["action"] == "INCORPORAR":
            helper_case = {
                "rut": case["rut"],
                "student": case["student"],
                "action": case["action"],
                "codcli": case["codcli"],
                "offer": case["offer"],
            }
            record, helper_trace = rect_helper.build_incorporation(helper_case, datos, h1, detalle_df)
        else:
            record, helper_trace, cesar_diff = rect_helper.build_cesar_annulment(headers)
        if rect_helper.offer_key(
            record["COD_SED"],
            record["COD_CAR"],
            record["MODALIDAD"],
            record["JOR"],
            record["VERSION"],
        ) != case["offer"]:
            fail(f"Oferta construida no coincide para {case['rut']}")
        if clean(record["N_DOC"]) != case["rut"]:
            fail(f"RUT construido no coincide para {case['rut']}")
        if clean(record["VIG"]) != case["vig"]:
            fail(f"VIG construido no coincide para {case['rut']}")
        records.append({col: clean(record[col]) for col in headers})
        trace_rows.append(
            {
                "N°": str(idx),
                "RUT": case["rut"],
                "Estudiante": case["student"],
                "Acción": case["action"],
                "CODCLI o registro": case["codcli_or_record"],
                "Oferta SIES": case["offer"],
                "VIG": case["vig"],
                "Fuente": clean(helper_trace.get("Fuente principal")),
                "Hoja": clean(helper_trace.get("Hoja fuente")),
                "Fila": clean(helper_trace.get("Fila fuente")),
                "Estado": "VALIDADO",
            }
        )

    out_df = pd.DataFrame(records, columns=headers).fillna("").astype(str)
    trace_df = pd.DataFrame(
        trace_rows,
        columns=["N°", "RUT", "Estudiante", "Acción", "CODCLI o registro", "Oferta SIES", "VIG", "Fuente", "Hoja", "Fila", "Estado"],
    ).fillna("").astype(str)

    included = set(out_df["N_DOC"].map(clean))
    expected = {case["rut"] for case in FINAL_CASES}
    if included != expected:
        fail(f"RUT incluidos no coinciden: {included} != {expected}")
    forbidden_included = included & (set(no_informar_ruts) | {MARCOS_RUT})
    if forbidden_included:
        fail(f"Se incluyeron RUT no accionables: {sorted(forbidden_included)}")
    return out_df, trace_df, cesar_diff


def write_rectification_csv(df: pd.DataFrame, template: Any) -> None:
    df.to_csv(
        RECTIFICATION_CSV,
        sep=template.delimiter,
        header=False if not template.has_header else True,
        index=False,
        encoding=template.encoding,
        lineterminator="\n",
        quoting=csv.QUOTE_MINIMAL,
    )


def base_header_style(cell: Any) -> None:
    cell.fill = PatternFill("solid", fgColor="D9EAF7")
    cell.font = Font(bold=True, color="1F4E79")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="thin", color="B7C9D6"))


def set_ws_common(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False


def write_rectification_excel(df: pd.DataFrame, trace_df: pd.DataFrame, headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RECTIFICACION_SIES"
    trace_ws = wb.create_sheet("TRAZABILIDAD")

    vig_fill = PatternFill("solid", fgColor="FFF2CC")
    cesar_fill = PatternFill("solid", fgColor="FCE4D6")

    for c_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        base_header_style(cell)
    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        is_cesar = clean(row["N_DOC"]) == "18059242" and clean(row["VIG"]) == "0"
        for c_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=clean(row[header]))
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if is_cesar:
                cell.fill = cesar_fill
            if header == "VIG":
                cell.fill = vig_fill
                cell.font = Font(bold=True, color="9C6500")
    set_ws_common(ws)
    for idx, header in enumerate(headers, start=1):
        max_data = max([len(clean(header))] + [len(clean(row[header])) for _, row in df.iterrows()])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_data + 2, 9), 24)

    trace_headers = list(trace_df.columns)
    for c_idx, header in enumerate(trace_headers, start=1):
        cell = trace_ws.cell(row=1, column=c_idx, value=header)
        base_header_style(cell)
    for r_idx, (_, row) in enumerate(trace_df.iterrows(), start=2):
        for c_idx, header in enumerate(trace_headers, start=1):
            cell = trace_ws.cell(row=r_idx, column=c_idx, value=clean(row[header]))
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    set_ws_common(trace_ws)
    widths = {
        "A": 6, "B": 12, "C": 34, "D": 18, "E": 42, "F": 16, "G": 8,
        "H": 58, "I": 36, "J": 38, "K": 14,
    }
    for col, width in widths.items():
        trace_ws.column_dimensions[col].width = width

    wb.save(RECTIFICATION_XLSX)


def append_table(ws: Any, headers: list[str], rows: list[list[Any]], start_row: int = 1, start_col: int = 1) -> None:
    for c_idx, header in enumerate(headers, start=start_col):
        cell = ws.cell(row=start_row, column=c_idx, value=header)
        base_header_style(cell)
    for r_offset, row in enumerate(rows, start=1):
        for c_offset, value in enumerate(row, start=0):
            cell = ws.cell(row=start_row + r_offset, column=start_col + c_offset, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_summary_excel() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN_EJECUTIVO"
    rect_ws = wb.create_sheet("RECTIFICACIONES")
    control_ws = wb.create_sheet("CONTROL_ENVIO")

    dark_blue = "1F4E79"
    light_blue = "D9EAF7"
    state_fill = PatternFill("solid", fgColor="E2F0D9")

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.35, bottom=0.35)
    ws.merge_cells("A1:H1")
    ws["A1"] = "Resumen Ejecutivo – Rectificación Matrícula Unificada 2026"
    ws["A1"].font = Font(size=20, bold=True, color=dark_blue)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:H2")
    ws["A2"] = "Resultado de la revisión de estudiantes con múltiples registros académicos"
    ws["A2"].font = Font(size=12, italic=True, color=dark_blue)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws["A4"] = "Contexto"
    ws["A4"].font = Font(bold=True, color=dark_blue)
    ws.merge_cells("A5:H6")
    ws["A5"] = (
        "Se revisó un universo de 34 estudiantes con múltiples CODCLI o trayectorias académicas. "
        "Como resultado del análisis, se identificaron cinco rectificaciones que requieren aprobación "
        "y posterior carga en SIES/PES."
    )
    ws["A5"].alignment = Alignment(wrap_text=True, vertical="top")

    append_table(
        ws,
        ["Indicador", "Resultado"],
        [
            ["Estudiantes revisados", 34],
            ["Rectificaciones propuestas", 5],
            ["Incorporaciones", 4],
            ["Anulaciones", 1],
        ],
        start_row=8,
        start_col=1,
    )
    for row in range(9, 13):
        ws.cell(row=row, column=2).font = Font(bold=True, size=12)

    ws["D8"] = "Decisión solicitada"
    ws["D8"].fill = PatternFill("solid", fgColor=light_blue)
    ws["D8"].font = Font(bold=True, color=dark_blue)
    ws.merge_cells("D9:H10")
    ws["D9"] = (
        "Se solicita revisar y aprobar las cinco rectificaciones propuestas. "
        "Una vez aprobadas, el archivo CSV adjunto podrá ser utilizado para la carga en SIES/PES."
    )
    ws["D9"].alignment = Alignment(wrap_text=True, vertical="top")

    ws["A14"] = "Archivos asociados"
    ws["A14"].font = Font(bold=True, color=dark_blue)
    files = [
        ("CSV", "archivo operativo de carga"),
        ("Excel con encabezados", "archivo para revisión y confirmación"),
        ("Auditable", "respaldo técnico interno"),
    ]
    append_table(ws, ["Archivo", "Uso"], files, start_row=15, start_col=1)

    ws.merge_cells("D15:H17")
    ws["D15"] = "PROCESO VALIDADO Y LISTO PARA REVISIÓN DE RECTORÍA"
    ws["D15"].fill = state_fill
    ws["D15"].font = Font(size=14, bold=True, color="375623")
    ws["D15"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.print_area = "A1:H18"
    for col, width in {"A": 28, "B": 18, "C": 4, "D": 24, "E": 18, "F": 18, "G": 18, "H": 18}.items():
        ws.column_dimensions[col].width = width
    for row in range(1, 19):
        ws.row_dimensions[row].height = 22

    rect_headers = ["N°", "RUT", "Estudiante", "Acción", "CODCLI o registro", "Oferta SIES", "VIG", "DECISIÓN RECTORÍA", "OBSERVACIÓN"]
    rect_rows = [
        [idx, c["rut"], c["student"], c["action"], c["codcli_or_record"], c["offer"], c["vig"], "", ""]
        for idx, c in enumerate(FINAL_CASES, start=1)
    ]
    append_table(rect_ws, rect_headers, rect_rows)
    set_ws_common(rect_ws)
    dv = DataValidation(type="list", formula1='"APROBADO,REVISAR,NO APROBADO"', allow_blank=True)
    rect_ws.add_data_validation(dv)
    dv.add(f"H2:H{len(rect_rows) + 1}")
    for col, width in {
        "A": 6, "B": 12, "C": 34, "D": 18, "E": 42, "F": 16, "G": 8, "H": 20, "I": 32,
    }.items():
        rect_ws.column_dimensions[col].width = width

    control_rows = [
        ["Universo revisado", 34, "OK"],
        ["Rectificaciones incluidas", 5, "OK"],
        ["Incorporaciones", 4, "OK"],
        ["Anulaciones", 1, "OK"],
        ["Registros CSV", 5, "OK"],
        ["Columnas CSV", 32, "OK"],
        ["Diferencias CSV vs Excel", 0, "OK"],
        ["Duplicados", 0, "OK"],
        ["Registros ajenos incluidos", 0, "OK"],
        ["Caso excluido incluido", 0, "OK"],
    ]
    append_table(control_ws, ["Control", "Resultado", "Estado"], control_rows)
    control_ws["A13"] = "El paquete contiene exclusivamente los cinco registros que requieren rectificación."
    control_ws["A13"].font = Font(bold=True, color=dark_blue)
    control_ws.merge_cells("A13:C13")
    set_ws_common(control_ws)
    for col, width in {"A": 32, "B": 16, "C": 12}.items():
        control_ws.column_dimensions[col].width = width

    wb.save(SUMMARY_XLSX)


def read_rectification_csv(template: Any, headers: list[str]) -> pd.DataFrame:
    return pd.read_csv(
        RECTIFICATION_CSV,
        sep=template.delimiter,
        header=0 if template.has_header else None,
        names=None if template.has_header else headers,
        dtype=str,
        encoding=template.encoding,
        keep_default_na=False,
    ).fillna("")


def compare_csv_vs_excel(template: Any, headers: list[str]) -> list[dict[str, str]]:
    csv_df = read_rectification_csv(template, headers)
    excel_df = pd.read_excel(RECTIFICATION_XLSX, sheet_name="RECTIFICACION_SIES", dtype=str, keep_default_na=False).fillna("")
    if list(csv_df.columns) != headers:
        fail("El CSV de rectificación no conserva el orden oficial de columnas")
    if list(excel_df.columns) != headers:
        fail("El Excel de revisión no conserva los encabezados oficiales")
    diffs: list[dict[str, str]] = []
    if csv_df.shape != excel_df.shape:
        return [{"FILA": "SHAPE", "COLUMNA": "SHAPE", "CSV": str(csv_df.shape), "EXCEL": str(excel_df.shape)}]
    for r_idx in range(len(csv_df)):
        for col in headers:
            csv_val = clean(csv_df.iloc[r_idx][col])
            excel_val = clean(excel_df.iloc[r_idx][col])
            if csv_val != excel_val:
                diffs.append({"FILA": str(r_idx + 1), "COLUMNA": col, "CSV": csv_val, "EXCEL": excel_val})
    return diffs


def extract_xlsx_values_text(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=False)
    pieces = list(wb.sheetnames)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    pieces.append(clean(value))
    wb.close()
    return "\n".join(pieces)


def validate_exclusions(no_informar_ruts: set[str]) -> dict[str, Any]:
    documents_text = {
        "resumen": extract_xlsx_values_text(SUMMARY_XLSX),
        "excel_revision": extract_xlsx_values_text(RECTIFICATION_XLSX),
        "csv": RECTIFICATION_CSV.read_text(encoding="utf-8"),
    }
    bad: list[dict[str, str]] = []
    for doc_name, text in documents_text.items():
        upper = text.upper()
        for token in FORBIDDEN_EXECUTIVE_TOKENS:
            if token.upper() in upper:
                bad.append({"DOCUMENTO": doc_name, "TOKEN": token})
        for rut in sorted(no_informar_ruts):
            if rut and rut in text:
                bad.append({"DOCUMENTO": doc_name, "TOKEN": rut})
    if bad:
        fail(f"Exclusiones fallidas en documentos ejecutivos/operativos: {bad[:10]}")
    return {"forbidden_hits": 0, "documents_checked": list(documents_text)}


def validate_rectification_outputs(rect_helper: Any, template: Any, headers: list[str], trace_df: pd.DataFrame, no_informar_ruts: set[str]) -> dict[str, Any]:
    csv_df = read_rectification_csv(template, headers)
    if csv_df.shape != (5, 32):
        fail(f"CSV rectificación esperado 5x32, obtenido {csv_df.shape}")
    if list(csv_df.columns) != headers:
        fail("CSV rectificación no conserva orden oficial")
    if RECTIFICATION_CSV.read_text(encoding=template.encoding).splitlines()[0].split(template.delimiter) == headers:
        fail("CSV rectificación contiene encabezados y el original no los tiene")
    if csv_df.apply(lambda col: col.astype(str).str.contains(r"E\+|E-", regex=True, na=False)).any().any():
        fail("Se detectó notación científica en CSV")

    wb = load_workbook(RECTIFICATION_XLSX, read_only=True, data_only=False)
    if wb.sheetnames != ["RECTIFICACION_SIES", "TRAZABILIDAD"]:
        fail(f"Excel revisión no tiene exactamente dos hojas: {wb.sheetnames}")
    ws = wb["RECTIFICACION_SIES"]
    if ws.max_row != 6 or ws.max_column != 32:
        fail(f"RECTIFICACION_SIES dimensión inválida: {ws.max_row}x{ws.max_column}")
    formulas = [
        cell.coordinate
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if formulas:
        fail(f"RECTIFICACION_SIES contiene fórmulas: {formulas[:5]}")
    wb.close()

    diffs = compare_csv_vs_excel(template, headers)
    if diffs:
        fail(f"Diferencias CSV vs Excel: {diffs[:10]}")

    exact_dups = int(csv_df.duplicated().sum())
    key_cols = ["TIPO_DOC", "N_DOC", "DV", "COD_SED", "COD_CAR", "MODALIDAD", "JOR", "VERSION"]
    key_dups = int(csv_df.duplicated(subset=key_cols).sum())
    if exact_dups or key_dups:
        fail(f"Duplicados detectados: exactos={exact_dups}, llave={key_dups}")

    included = set(csv_df["N_DOC"].map(clean))
    expected = {c["rut"] for c in FINAL_CASES}
    if included != expected:
        fail(f"RUT incluidos no coinciden: {included} != {expected}")
    alien = sorted(included & (set(no_informar_ruts) | {MARCOS_RUT}))
    if alien:
        fail(f"Registros ajenos incluidos: {alien}")

    final_case_rows = []
    for case in FINAL_CASES:
        rows = csv_df[csv_df["N_DOC"].map(clean).eq(case["rut"])]
        if len(rows) != 1:
            fail(f"{case['rut']}: se esperaba una fila y hay {len(rows)}")
        row = rows.iloc[0]
        offer = rect_helper.offer_key(row["COD_SED"], row["COD_CAR"], row["MODALIDAD"], row["JOR"], row["VERSION"])
        if offer != case["offer"] or clean(row["VIG"]) != case["vig"]:
            fail(f"{case['rut']}: oferta/VIG no coincide: {offer}, VIG={row['VIG']}")
        if case["rut"] == "18059242":
            for col in ["PROM_PRI_SEM", "PROM_SEG_SEM", "ASI_INS_HIS", "ASI_APR_HIS"]:
                if clean(row[col]) != "0":
                    fail(f"César no tiene {col}=0")
        final_case_rows.append(
            {
                "RUT": case["rut"],
                "Estudiante": case["student"],
                "Acción": case["action"],
                "CODCLI o registro": case["codcli_or_record"],
                "Oferta": case["offer"],
                "VIG": case["vig"],
                "Estado": "VALIDADO",
            }
        )

    wb_summary = load_workbook(SUMMARY_XLSX, read_only=True, data_only=False)
    if wb_summary.sheetnames != ["RESUMEN_EJECUTIVO", "RECTIFICACIONES", "CONTROL_ENVIO"]:
        fail(f"Resumen ejecutivo no tiene exactamente 3 hojas: {wb_summary.sheetnames}")
    wb_summary.close()
    validate_exclusions(no_informar_ruts)
    ext_review = validate_no_external_links(RECTIFICATION_XLSX)
    ext_summary = validate_no_external_links(SUMMARY_XLSX)
    if ext_review != (0, 0) or ext_summary != (0, 0):
        fail(f"Vínculos externos detectados: review={ext_review}, summary={ext_summary}")

    return {
        "csv_rows": len(csv_df),
        "csv_cols": csv_df.shape[1],
        "csv_excel_diffs": len(diffs),
        "exact_dups": exact_dups,
        "key_dups": key_dups,
        "alien_records": 0,
        "rut_15651488_included": 0,
        "vig0_count": int((csv_df["VIG"].map(clean) == "0").sum()),
        "incorporations": sum(1 for c in FINAL_CASES if c["action"] == "INCORPORAR"),
        "annulments": sum(1 for c in FINAL_CASES if c["action"] == "ELIMINAR/ANULAR"),
        "case_rows": final_case_rows,
    }


def write_email_body() -> None:
    body = """Asunto sugerido: Revisión y aprobación de rectificación Matrícula Unificada 2026

Estimados/as:

Junto con saludar, adjunto los antecedentes correspondientes a la revisión de Matrícula Unificada 2026.

El análisis consideró un universo de 34 estudiantes con múltiples CODCLI o trayectorias académicas. Como resultado, se identificaron cinco rectificaciones que requieren revisión y aprobación: cuatro incorporaciones y una anulación.

Para la revisión se adjuntan:

1. RESUMEN_EJECUTIVO_CIERRE_34_MULTIVIGENTES_SIES.xlsx, que presenta el resultado consolidado y las cinco acciones propuestas.
2. RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026_CON_ENCABEZADOS.xlsx, que permite revisar los registros completos con sus títulos de columna.
3. CIERRE_34_MULTIVIGENTES_SIES_AUDITABLE.xlsx, que contiene el respaldo técnico y la trazabilidad del análisis.
4. RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026.csv, que corresponde al archivo operativo preparado con la estructura oficial de carga y que deberá utilizarse en SIES/PES una vez aprobadas las rectificaciones.

El archivo CSV contiene exclusivamente los cinco registros que requieren gestión y mantiene la misma estructura de 32 columnas utilizada en la carga original. El Excel con encabezados contiene los mismos valores del CSV y se incorpora para facilitar su revisión.

Agradeceré revisar y confirmar las cinco rectificaciones propuestas. Una vez obtenida la aprobación, el archivo CSV quedará disponible para su carga en SIES/PES.

Saludos cordiales,
"""
    EMAIL_BODY_TXT.write_text(body, encoding="utf-8")


def validate_email_body() -> None:
    text = EMAIL_BODY_TXT.read_text(encoding="utf-8")
    required = [
        "universo de 34 estudiantes",
        "cinco rectificaciones",
        "cuatro incorporaciones",
        "una anulación",
        "RESUMEN_EJECUTIVO_CIERRE_34_MULTIVIGENTES_SIES.xlsx",
        "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026_CON_ENCABEZADOS.xlsx",
        "CIERRE_34_MULTIVIGENTES_SIES_AUDITABLE.xlsx",
        "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026.csv",
        "SIES/PES",
        "Excel con encabezados",
    ]
    lower = text.lower()
    missing = [phrase for phrase in required if phrase.lower() not in lower]
    if missing:
        fail(f"Cuerpo de correo no contiene textos requeridos: {missing}")
    forbidden = ["28", "Marcos", "NO_DETERMINABLE", "NO_INFORMAR", MARCOS_RUT]
    hits = [token for token in forbidden if token.lower() in lower]
    if hits:
        fail(f"Cuerpo de correo contiene textos prohibidos: {hits}")


def copy_with_backup(src: Path, dst: Path, backups: list[Path]) -> None:
    backup_if_exists(dst, backups)
    shutil.copy2(src, dst)


def print_table(rows: list[dict[str, Any]], columns: list[str]) -> None:
    widths = {
        col: max(len(col), *(len(clean(row.get(col, ""))) for row in rows))
        for col in columns
    }
    print(" | ".join(col.ljust(widths[col]) for col in columns))
    print(" | ".join("-" * widths[col] for col in columns))
    for row in rows:
        print(" | ".join(clean(row.get(col, "")).ljust(widths[col]) for col in columns))


def validate_summary_workbook() -> dict[str, Any]:
    wb = load_workbook(SUMMARY_XLSX, read_only=True, data_only=True)
    if wb.sheetnames != ["RESUMEN_EJECUTIVO", "RECTIFICACIONES", "CONTROL_ENVIO"]:
        fail(f"Resumen ejecutivo hojas inválidas: {wb.sheetnames}")
    rect_rows = rows_from_ws(wb["RECTIFICACIONES"])
    if len(rect_rows) != 5:
        fail(f"Resumen ejecutivo RECTIFICACIONES esperado 5 filas, obtenido {len(rect_rows)}")
    control_rows = rows_from_ws(wb["CONTROL_ENVIO"])
    bad = [r for r in control_rows if clean(r.get("Estado")) and clean(r.get("Estado")) != "OK"]
    if bad:
        fail(f"Resumen ejecutivo controles no OK: {bad}")
    wb.close()
    return {"sheets": 3, "rectification_rows": len(rect_rows)}


def document_report_row(name: str, project: Path, desktop: Path, rows: str, cols_or_sheets: str) -> dict[str, str]:
    return {
        "Documento": name,
        "Ruta proyecto": str(project),
        "Ruta Escritorio": str(desktop),
        "Filas": rows,
        "Columnas u hojas": cols_or_sheets,
        "SHA-256 idéntico": "SI" if sha256(project) == sha256(desktop) else "NO",
    }


def main() -> int:
    if BASE.resolve() != Path.cwd().resolve():
        print(f"Directorio actual: {Path.cwd()}")
        print(f"Cambiando contexto lógico a: {BASE}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    backups: list[Path] = []
    rect_helper = import_rectification_helper()

    # Fase 1: respaldos y hashes previos.
    final_project_files = [SUMMARY_XLSX, RECTIFICATION_CSV, RECTIFICATION_XLSX, AUDITABLE_XLSX, EMAIL_BODY_TXT]
    final_desktop_files = [
        DESKTOP_SUMMARY_XLSX,
        DESKTOP_RECTIFICATION_CSV,
        DESKTOP_RECTIFICATION_XLSX,
        DESKTOP_AUDITABLE_XLSX,
        DESKTOP_EMAIL_BODY_TXT,
    ]
    for path in final_project_files + final_desktop_files:
        backup_if_exists(path, backups)
    if MASTER_SCRIPT.exists() and MASTER_SCRIPT.resolve() != Path(__file__).resolve():
        backup_if_exists(MASTER_SCRIPT, backups)
    source_hashes_before = snapshot_hashes(source_paths())
    technical_hash_before = sha256(TECHNICAL_XLSX)

    print("RESPALDOS_CREADOS")
    if backups:
        for backup in backups:
            print(backup)
    else:
        print("No existían archivos finales o scripts previos que respaldar.")
    print("\nHASHES_FUENTES_ANTES")
    for path, digest in source_hashes_before.items():
        print(f"{digest}  {path}")

    headers = rect_helper.load_official_headers()
    if headers != MU32_COLUMNS:
        fail("Encabezados oficiales no coinciden con el orden validado de 32 columnas")
    candidates = rect_helper.discover_csv_candidates(headers)
    template = rect_helper.select_template(candidates)
    original_df = rect_helper.load_original_dataframe(template, headers)
    if template.path != OFFICIAL_CSV:
        fail(f"CSV oficial seleccionado inesperado: {template.path}")
    if len(original_df) != 4070 or original_df.shape[1] != 32:
        fail(f"CSV oficial debe ser 4070x32, obtenido {original_df.shape}")

    print("\nCSV_ORIGINAL_OFICIAL")
    print(f"ruta: {template.path}")
    print(f"filas: {len(original_df)}")
    print(f"columnas: {original_df.shape[1]}")
    print(f"delimitador: {repr(template.delimiter)}")
    print(f"codificación: {template.encoding}")
    print(f"encabezados: {'SI' if template.has_header else 'NO'}")
    print("orden de columnas:")
    for idx, col in enumerate(headers, start=1):
        print(f"{idx:02d}. {col}")

    # Fase 2: auditable.
    run_checked([sys.executable, str(AUDITABLE_SCRIPT)], label="regenerar auditable técnico")
    recalc_with_microsoft_excel(AUDITABLE_XLSX)
    auditable_info = validate_auditable(rect_helper)

    # Fases 3 y 4: CSV y Excel con encabezados.
    out_df, trace_df, cesar_diff = build_rectification_records(rect_helper, template, headers)
    write_rectification_csv(out_df, template)
    write_rectification_excel(out_df, trace_df, headers)

    # Fase 5: resumen ejecutivo.
    write_summary_excel()
    summary_info = validate_summary_workbook()

    # Fases 6, 7 y 8: validaciones de contenido y exclusión.
    validation = validate_rectification_outputs(
        rect_helper,
        template,
        headers,
        trace_df,
        auditable_info["no_informar_ruts"],
    )

    # Fase 12 y 13: cuerpo de correo.
    write_email_body()
    validate_email_body()

    # Fase 9: copias al Escritorio.
    desktop_backups: list[Path] = []
    copy_with_backup(SUMMARY_XLSX, DESKTOP_SUMMARY_XLSX, desktop_backups)
    copy_with_backup(RECTIFICATION_CSV, DESKTOP_RECTIFICATION_CSV, desktop_backups)
    copy_with_backup(RECTIFICATION_XLSX, DESKTOP_RECTIFICATION_XLSX, desktop_backups)
    copy_with_backup(AUDITABLE_XLSX, DESKTOP_AUDITABLE_XLSX, desktop_backups)
    copy_with_backup(EMAIL_BODY_TXT, DESKTOP_EMAIL_BODY_TXT, desktop_backups)
    backups.extend(desktop_backups)

    # Fase 10: hashes e integridad.
    output_pairs = [
        (SUMMARY_XLSX, DESKTOP_SUMMARY_XLSX),
        (RECTIFICATION_CSV, DESKTOP_RECTIFICATION_CSV),
        (RECTIFICATION_XLSX, DESKTOP_RECTIFICATION_XLSX),
        (AUDITABLE_XLSX, DESKTOP_AUDITABLE_XLSX),
        (EMAIL_BODY_TXT, DESKTOP_EMAIL_BODY_TXT),
    ]
    hash_pairs = [(project, desktop, sha256(project), sha256(desktop)) for project, desktop in output_pairs]
    mismatches = [(p, d) for p, d, hp, hd in hash_pairs if hp != hd]
    if mismatches:
        fail(f"Hashes proyecto/Escritorio no coinciden: {mismatches}")

    source_hashes_after = snapshot_hashes(source_paths())
    changed_sources = [
        path for path in source_hashes_before
        if source_hashes_before[path] != source_hashes_after[path]
    ]
    if changed_sources:
        fail(f"Fuentes originales modificadas: {changed_sources}")
    if sha256(TECHNICAL_XLSX) != technical_hash_before:
        fail("El cierre técnico fue modificado")

    print("\nCOMPARACION_CESAR_ANTES_DESPUES")
    if cesar_diff:
        print_table(cesar_diff, ["COLUMNA", "ANTES", "DESPUES", "CAMBIA"])

    print("\nHASHES_SALIDAS")
    for project, desktop, project_hash, desktop_hash in hash_pairs:
        print(f"{project_hash}  {project}")
        print(f"{desktop_hash}  {desktop}")
    print("\nHASHES_FUENTES_DESPUES")
    for path, digest in source_hashes_after.items():
        print(f"{digest}  {path}")

    print("\nDOCUMENTOS_GENERADOS")
    doc_rows = [
        document_report_row("Resumen ejecutivo", SUMMARY_XLSX, DESKTOP_SUMMARY_XLSX, "3 hojas", "3 hojas"),
        document_report_row("CSV rectificación", RECTIFICATION_CSV, DESKTOP_RECTIFICATION_CSV, "5", "32"),
        document_report_row("Excel con encabezados", RECTIFICATION_XLSX, DESKTOP_RECTIFICATION_XLSX, "5 en RECTIFICACION_SIES", "2 hojas"),
        document_report_row("Auditable técnico", AUDITABLE_XLSX, DESKTOP_AUDITABLE_XLSX, "34 RUT / 68 CODCLI", "10 hojas"),
        document_report_row("Cuerpo correo", EMAIL_BODY_TXT, DESKTOP_EMAIL_BODY_TXT, "texto", "1 archivo"),
    ]
    print_table(doc_rows, ["Documento", "Ruta proyecto", "Ruta Escritorio", "Filas", "Columnas u hojas", "SHA-256 idéntico"])

    print("\nRESUMEN_CONTROLES")
    print(f"universo revisado: {auditable_info['rut_total']}")
    print(f"rectificaciones: {len(FINAL_CASES)}")
    print(f"incorporaciones: {validation['incorporations']}")
    print(f"anulaciones: {validation['annulments']}")
    print(f"registros CSV: {validation['csv_rows']}")
    print(f"columnas CSV: {validation['csv_cols']}")
    print(f"diferencias CSV vs Excel: {validation['csv_excel_diffs']}")
    print(f"duplicados: {validation['exact_dups'] + validation['key_dups']}")
    print(f"registros ajenos incluidos: {validation['alien_records']}")
    print(f"RUT 15651488 en documentos ejecutivos y operativos: {validation['rut_15651488_included']}")
    print("fuentes originales modificadas: NO")
    print("cuerpo de correo generado: SÍ")
    print(f"auditable fórmulas totales: {auditable_info['total_formulas']}")
    print(f"auditable ES_2026_1: {auditable_info['es_2026']}")
    print(f"auditable clasificaciones: {auditable_info['class_counts']}")
    print(f"diferencias reales contra cierre técnico: {len(auditable_info['differences_against_technical'])}")
    print(f"vínculos externos auditable: {auditable_info['external_links']}")
    print(f"fórmulas externas auditable: {auditable_info['external_formulas']}")
    print(f"hojas resumen ejecutivo: {summary_info['sheets']}")

    print("\nCUERPO_CORREO")
    print(EMAIL_BODY_TXT.read_text(encoding="utf-8"))

    print("\nTABLA_EJECUTIVA_FINAL")
    print_table(validation["case_rows"], ["RUT", "Estudiante", "Acción", "CODCLI o registro", "Oferta", "VIG", "Estado"])

    print("\nOK: LOS CUATRO DOCUMENTOS FUERON REGENERADOS Y VALIDADOS.")
    print("OK: EL CSV CONTIENE EXCLUSIVAMENTE LOS CINCO REGISTROS DE RECTIFICACIÓN.")
    print("OK: EL EXCEL CON ENCABEZADOS ES IDÉNTICO AL CSV.")
    print("OK: EL RESUMEN EJECUTIVO SOLO PRESENTA EL UNIVERSO DE 34 ESTUDIANTES Y LAS CINCO ACCIONES.")
    print("OK: EL RUT 15651488 Y LOS CASOS SIN ACCIÓN FUERON EXCLUIDOS DE LOS DOCUMENTOS EJECUTIVOS Y OPERATIVOS.")
    print("OK: EL ARCHIVO AUDITABLE CONSERVA LA TRAZABILIDAD TÉCNICA COMPLETA.")
    print("OK: EL CUERPO DEL CORREO FUE GENERADO Y ESTÁ LISTO PARA REVISIÓN.")
    print("OK: PAQUETE FINAL LISTO PARA ENVIAR A RECTORÍA.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
