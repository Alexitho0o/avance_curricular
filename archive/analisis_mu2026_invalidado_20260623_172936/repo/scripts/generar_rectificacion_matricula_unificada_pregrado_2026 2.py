#!/usr/bin/env python3
from __future__ import annotations

import csv
import hashlib
import re
import shutil
import sys
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


BASE = Path("/Users/alexi/Documents/GitHub/avance_curricular")
OUT_DIR = BASE / "resultados/auditoria_multicodcli_reconstruida_2026/cierre_34_multivigentes_sies_20260623"
DESKTOP = Path("/Users/alexi/Desktop")

AUDITABLE_XLSX = OUT_DIR / "CIERRE_34_MULTIVIGENTES_SIES_AUDITABLE.xlsx"
TECNICO_XLSX = OUT_DIR / "CIERRE_34_MULTIVIGENTES_SIES.xlsx"
PROMEDIOS_XLSX = BASE / "input/PROMEDIOSDEALUMNOS_7804.xlsx"
TEMPLATE_CSV = BASE / "control/punto_0_carga_principal_mu2026/archivos_congelados/PARA_SUBIR_DESKTOP__matricula_unificada_2026_pregrado_PARA_SUBIR.csv"
PES_READY_CSV = BASE / "resultados/matricula_unificada_2026_pregrado_PES_READY.csv"
HEADER_CSV = BASE / "resultados/matricula_unificada_2026_control.csv"
CESAR_SOURCE_CSV = BASE / "resultados/matricula_unificada_rectificacion_20260618/matricula_unificada_2026_BASE_COMPLETA_CORREGIDA_P1_20260618.csv"

OUT_CSV = OUT_DIR / "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026.csv"
OUT_XLSX = OUT_DIR / "RECTIFICACION_MATRICULA_UNIFICADA_PREGRADO_2026_CON_ENCABEZADOS.xlsx"
DESKTOP_CSV = DESKTOP / OUT_CSV.name
DESKTOP_XLSX = DESKTOP / OUT_XLSX.name

EXPECTED_UPLOAD_HASH = "8ac3d58613d000534bf4053a5b9e42d13eee5db488f00d62d3f6d035c4df2704"

MU32_COLUMNS = [
    "TIPO_DOC", "N_DOC", "DV", "PRIMER_APELLIDO", "SEGUNDO_APELLIDO", "NOMBRE",
    "SEXO", "FECH_NAC", "NAC", "PAIS_EST_SEC", "COD_SED", "COD_CAR", "MODALIDAD",
    "JOR", "VERSION", "FOR_ING_ACT", "ANIO_ING_ACT", "SEM_ING_ACT", "ANIO_ING_ORI",
    "SEM_ING_ORI", "ASI_INS_ANT", "ASI_APR_ANT", "PROM_PRI_SEM", "PROM_SEG_SEM",
    "ASI_INS_HIS", "ASI_APR_HIS", "NIV_ACA", "SIT_FON_SOL", "SUS_PRE",
    "FECHA_MATRICULA", "REINCORPORACION", "VIG",
]

RECTIFICATION_CASES = [
    {
        "rut": "17726298",
        "student": "JOSE JAIRO VELASQUEZ FIGUEROA",
        "action": "INCORPORAR",
        "codcli": "20253IINF012",
        "offer": "S2C1M3J4V2",
    },
    {
        "rut": "18939583",
        "student": "ADOLFO ANDRES CAMPOS GOMEZ",
        "action": "INCORPORAR",
        "codcli": "20241TPAS250",
        "offer": "S2C57M3J4V1",
    },
    {
        "rut": "19356713",
        "student": "NICOLAS OCTAVIO LAGOS VEGA",
        "action": "INCORPORAR",
        "codcli": "20251ICDA004",
        "offer": "S2C88M3J4V1",
    },
    {
        "rut": "19513133",
        "student": "ARIEL EDUARDO MARTINEZ ALVAREZ",
        "action": "INCORPORAR",
        "codcli": "20261ICDA007",
        "offer": "S2C88M3J4V1",
    },
    {
        "rut": "18059242",
        "student": "CESAR ANDRES RUBILAR SANHUEZA",
        "action": "ELIMINAR/ANULAR",
        "codcli": "",
        "registro": "MU_LINEA_04073|RUT=18059242-3|S2C1M1J1V1",
        "offer": "S2C1M1J1V1",
    },
]

MARCOS_RUT = "15651488"
MARCOS_CODCLI = "20261DDASC008"

NAC_MAP = {
    "ALEMANA": "3", "ARGENTINA": "9", "BOLIVIANA": "23", "BRASILENA": "26",
    "BRASILEÑA": "26", "CHILENA": "38", "COLOMBIANA": "44", "CUBANA": "48",
    "ECUATORIANA": "57", "ESPANOLA": "68", "ESPAÑOLA": "68", "FRANCESA": "70",
    "HAITIANA": "78", "ITALIANA": "95", "MEXICANA": "130", "PARAGUAYA": "137",
    "PERUANA": "142", "URUGUAYA": "188", "VENEZOLANA": "192",
}


@dataclass
class CandidateCsv:
    path: Path
    rows: int
    cols: int
    delimiter: str
    encoding: str
    has_header: bool
    sha256: str
    reason: str


def fail(message: str) -> None:
    raise RuntimeError(message)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat", "<na>"}:
        return ""
    if re.fullmatch(r"-?\d+\.0", text):
        return text[:-2]
    return re.sub(r"\s+", " ", text).strip()


def norm_text(value: Any, *, strip_accents: bool = True) -> str:
    text = clean(value).upper()
    if strip_accents:
        text = unicodedata.normalize("NFKD", text)
        text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text).strip()


def rut_norm(value: Any) -> str:
    text = clean(value).replace(".", "")
    if "-" in text:
        text = text.split("-", 1)[0]
    return re.sub(r"\D", "", text)


def parse_rut(value: Any) -> tuple[str, str]:
    text = clean(value).replace(".", "").upper()
    if "-" not in text:
        return re.sub(r"\D", "", text), ""
    n_doc, dv = text.rsplit("-", 1)
    return re.sub(r"\D", "", n_doc), re.sub(r"[^0-9K]", "", dv)


def date_to_slash(value: Any) -> str:
    text = clean(value)
    if not text:
        return ""
    if re.fullmatch(r"\d{2}-\d{2}-\d{4}", text):
        return text.replace("-", "/")
    if re.fullmatch(r"\d{2}/\d{2}/\d{4}", text):
        return text
    parsed = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return text
    return parsed.strftime("%d/%m/%Y")


def offer_key(sede: Any, cod_car: Any, modalidad: Any, jornada: Any, version: Any) -> str:
    vals = [clean(sede), clean(cod_car), clean(modalidad), clean(jornada), clean(version)]
    if not all(vals):
        return ""
    return f"S{vals[0]}C{vals[1]}M{vals[2]}J{vals[3]}V{vals[4]}"


def parse_offer(text: Any) -> tuple[str, str, str, str, str]:
    match = re.fullmatch(r"S(\d+)C(\d+)M(\d+)J(\d+)V(\d+)", clean(text).upper())
    if not match:
        fail(f"Oferta SIES inválida: {text!r}")
    return match.groups()


def period_to_semester(value: Any) -> str:
    text = clean(value)
    if text == "1":
        return "1"
    if text in {"2", "3"}:
        return "2"
    fail(f"PERIODOINGRESO fuera de catálogo {value!r}")


def level_to_semester(value: Any) -> str:
    text = clean(value)
    if not text:
        fail("NIVEL vacío para derivar NIV_ACA")
    level = int(float(text))
    mapping = {
        1: 1, 2: 2, 3: 2, 4: 3, 5: 4, 6: 4,
        7: 5, 8: 6, 9: 6, 10: 7, 11: 8, 12: 8,
    }
    if level in mapping:
        return str(mapping[level])
    return str((level * 2 + 2) // 3)


def normalize_grade_to_mu(value: Any) -> int | None:
    text = clean(value)
    if not text:
        return None
    number = pd.to_numeric(pd.Series([text]), errors="coerce").iloc[0]
    if pd.isna(number):
        return None
    number = float(number)
    if 1 <= number <= 7:
        return int(round(number * 100))
    if 7 < number <= 70:
        return int(round(number * 10))
    if 100 <= number <= 700:
        return int(round(number))
    return None


def coerce_mu_average(values: list[int]) -> str:
    values = [v for v in values if v is not None]
    if not values:
        return "0"
    avg = int(round(sum(values) / len(values)))
    if avg == 0:
        return "0"
    return str(min(max(avg, 100), 700))


def detect_encoding(path: Path) -> str:
    raw = path.read_bytes()[:4]
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            path.read_text(encoding=enc)
            return enc
        except UnicodeDecodeError:
            continue
    fail(f"No se pudo detectar codificación para {path}")


def read_csv_rows(path: Path, delimiter: str, encoding: str) -> list[list[str]]:
    with path.open("r", encoding=encoding, newline="") as f:
        return list(csv.reader(f, delimiter=delimiter))


def detect_delimiter(path: Path, encoding: str) -> str:
    sample = path.read_text(encoding=encoding, errors="replace")[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=[",", ";", "\t", "|"])
        return dialect.delimiter
    except csv.Error:
        first = sample.splitlines()[0]
        counts = {d: first.count(d) for d in [",", ";", "\t", "|"]}
        return max(counts, key=counts.get)


def has_header(first_row: list[str], headers: list[str]) -> bool:
    return [clean(x).upper() for x in first_row] == [h.upper() for h in headers]


def load_official_headers() -> list[str]:
    if not HEADER_CSV.exists():
        fail(f"No existe diccionario/control de encabezados: {HEADER_CSV}")
    df = pd.read_csv(HEADER_CSV, dtype=str, nrows=0, encoding="utf-8-sig")
    headers = list(df.columns)
    if headers != MU32_COLUMNS:
        fail(f"Encabezados del control no coinciden con MU32 esperado: {headers}")
    return headers


def discover_csv_candidates(headers: list[str]) -> list[CandidateCsv]:
    candidates: list[CandidateCsv] = []
    for path in BASE.rglob("*.csv"):
        try:
            enc = detect_encoding(path)
            delim = detect_delimiter(path, enc)
            rows = read_csv_rows(path, delim, enc)
        except Exception:
            continue
        if not rows:
            continue
        first_cols = len(rows[0])
        if first_cols != 32:
            continue
        header = has_header(rows[0], headers)
        data_rows = len(rows) - (1 if header else 0)
        rel = str(path.relative_to(BASE))
        reason = ""
        digest = sha256(path)
        if digest == EXPECTED_UPLOAD_HASH:
            reason = "HASH_CARGA_PRINCIPAL_VALIDADO"
        elif "PES_READY" in rel or "PARA_SUBIR" in rel:
            reason = "CANDIDATO_NOMBRE_CARGA"
        elif "matricula_unificada_2026_pregrado" in rel:
            reason = "CANDIDATO_PREGRADO_2026"
        else:
            reason = "CSV_32_COLUMNAS"
        candidates.append(CandidateCsv(path, data_rows, 32, delim, enc, header, digest, reason))
    candidates.sort(
        key=lambda c: (
            c.sha256 == EXPECTED_UPLOAD_HASH,
            "PARA_SUBIR_DESKTOP" in str(c.path),
            "PES_READY" in str(c.path),
            c.rows,
            c.path.stat().st_mtime,
        ),
        reverse=True,
    )
    return candidates


def select_template(candidates: list[CandidateCsv]) -> CandidateCsv:
    for candidate in candidates:
        if candidate.path == TEMPLATE_CSV:
            if candidate.sha256 != EXPECTED_UPLOAD_HASH:
                fail(f"El CSV congelado de carga no tiene el hash validado: {candidate.sha256}")
            return candidate
    fail(f"No se encontró el CSV oficial seleccionado: {TEMPLATE_CSV}")


def load_auditable() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    casos = pd.read_excel(AUDITABLE_XLSX, sheet_name="01_CASOS_TRAZABLES", dtype=str)
    detalle = pd.read_excel(AUDITABLE_XLSX, sheet_name="02_DETALLE_CODCLI", dtype=str)
    filas_mu = pd.read_excel(AUDITABLE_XLSX, sheet_name="04_FILAS_MU", dtype=str)
    for df in [casos, detalle, filas_mu]:
        df.columns = [clean(c) for c in df.columns]
    return casos, detalle, filas_mu


def validate_decision_sources(casos: pd.DataFrame, detalle: pd.DataFrame) -> tuple[set[str], set[str]]:
    class_counts = Counter(casos["CLASIFICACION_FINAL"].map(clean))
    if len(casos) != 34:
        fail(f"Auditable no tiene 34 RUT: {len(casos)}")
    expected = {"CORRECCION_SIES": 5, "NO_INFORMAR": 28, "NO_DETERMINABLE": 1}
    for key, value in expected.items():
        if class_counts.get(key, 0) != value:
            fail(f"Clasificación {key} esperada {value}, obtenida {class_counts.get(key, 0)}")
    if len(detalle) != 68:
        fail(f"Detalle CODCLI esperado 68, obtenido {len(detalle)}")
    if clean(detalle.loc[detalle["CODCLI"].map(clean).eq(MARCOS_CODCLI), "MAPEO_DEMOSTRADO"].iloc[0]) != "NO":
        fail("Marcos Quezada no figura como mapeo NO para 20261DDASC008")
    no_informar = set(casos.loc[casos["CLASIFICACION_FINAL"].map(clean).eq("NO_INFORMAR"), "RUT_NORM"].map(rut_norm))
    no_determinable = set(casos.loc[casos["CLASIFICACION_FINAL"].map(clean).eq("NO_DETERMINABLE"), "RUT_NORM"].map(rut_norm))
    return no_informar, no_determinable


def load_original_dataframe(template: CandidateCsv, headers: list[str]) -> pd.DataFrame:
    df = pd.read_csv(
        template.path,
        sep=template.delimiter,
        header=None if not template.has_header else 0,
        names=headers if not template.has_header else None,
        dtype=str,
        encoding=template.encoding,
        keep_default_na=False,
    )
    if list(df.columns) != headers:
        fail("El orden de columnas del CSV original no coincide con el diccionario validado")
    if df.shape[1] != 32:
        fail(f"CSV original tiene {df.shape[1]} columnas")
    return df.fillna("")


def get_unique_row(df: pd.DataFrame, col: str, value: str, label: str) -> pd.Series:
    matches = df[df[col].map(clean).eq(value)]
    if len(matches) != 1:
        fail(f"{label}: se esperaba 1 fila para {col}={value}, hay {len(matches)}")
    return matches.iloc[0]


def compute_historical_fields(h1: pd.DataFrame, codcli: str, anio_ref: int = 2025) -> tuple[dict[str, str], str]:
    sub = h1[h1["CODCLI"].map(clean).eq(codcli)].copy()
    if sub.empty:
        return {
            "ASI_INS_ANT": "0", "ASI_APR_ANT": "0", "PROM_PRI_SEM": "0",
            "PROM_SEG_SEM": "0", "ASI_INS_HIS": "0", "ASI_APR_HIS": "0",
        }, "Hoja1 sin filas para CODCLI; campos historicos en 0"

    sub["ANO_NUM"] = pd.to_numeric(sub["ANO"], errors="coerce")
    sub["SEMESTRE_HIST"] = sub["PERIODO"].map(period_to_semester)
    sub["NOTA_MU"] = sub["NOTA_FINAL"].map(normalize_grade_to_mu)
    sub["ESTADO_HIST_NORM"] = sub["DESCRIPCION_ESTADO"].map(norm_text)
    sub["CONVALIDADO_NORM"] = sub.get("CONVALIDADO", pd.Series(index=sub.index, dtype=str)).map(norm_text)
    sub_ref = sub[sub["ANO_NUM"].eq(anio_ref)].copy()
    estado_ref = sub_ref["ESTADO_HIST_NORM"]
    transfer_ref = (
        estado_ref.str.contains(r"CONVALID|HOMOLOG|RECONOC|EQUIV", regex=True, na=False)
        | sub_ref["CONVALIDADO_NORM"].eq("S")
    )
    graded_ref = sub_ref["NOTA_MU"].notna() & ~transfer_ref
    aprob_ref = estado_ref.str.contains("APROB", na=False) & ~transfer_ref
    aprob_hist = sub["ESTADO_HIST_NORM"].str.contains(r"APROB|CONVALID|RECONOC|EQUIV|HOMOLOG", regex=True, na=False)

    fields = {
        "ASI_INS_ANT": str(int(sub_ref.loc[~transfer_ref, "CODRAMO"].nunique())),
        "ASI_APR_ANT": str(int(sub_ref.loc[aprob_ref, "CODRAMO"].nunique())),
        "PROM_PRI_SEM": coerce_mu_average(sub_ref.loc[graded_ref & sub_ref["SEMESTRE_HIST"].eq("1"), "NOTA_MU"].dropna().astype(int).tolist()),
        "PROM_SEG_SEM": coerce_mu_average(sub_ref.loc[graded_ref & sub_ref["SEMESTRE_HIST"].eq("2"), "NOTA_MU"].dropna().astype(int).tolist()),
        "ASI_INS_HIS": str(int(sub["CODRAMO"].dropna().count())),
        "ASI_APR_HIS": str(int(sub.loc[aprob_hist, "CODRAMO"].nunique())),
    }
    evidence = (
        f"Hoja1 CODCLI {codcli}: filas total={len(sub)}, filas {anio_ref}={len(sub_ref)}, "
        f"transferencias_ref={int(transfer_ref.sum())}"
    )
    return fields, evidence


def build_incorporation(
    case: dict[str, str],
    datos: pd.DataFrame,
    h1: pd.DataFrame,
    detalle: pd.DataFrame,
) -> tuple[dict[str, str], dict[str, str]]:
    codcli = case["codcli"]
    rut = case["rut"]
    expected_offer = case["offer"]

    da = get_unique_row(datos, "CODCLI", codcli, f"DatosAlumnos {codcli}")
    da_excel_row = str(int(da.name) + 2)
    n_doc, dv = parse_rut(da["RUT"])
    if n_doc != rut:
        fail(f"{codcli}: RUT fuente {n_doc} no coincide con {rut}")
    if clean(da["ESTADOACADEMICO"]).upper() != "VIGENTE":
        fail(f"{codcli}: ESTADOACADEMICO no vigente")
    if not clean(da["MATRICULA"]):
        fail(f"{codcli}: MATRICULA vacía")

    detail = detalle[(detalle["CODCLI"].map(clean).eq(codcli)) & (detalle["RUT_NORM"].map(rut_norm).eq(rut))]
    if len(detail) != 1:
        fail(f"{codcli}: no hay detalle auditable único")
    detail_row = detail.iloc[0]
    if clean(detail_row["MAPEO_DEMOSTRADO"]) != "SI":
        fail(f"{codcli}: mapeo no demostrado")
    if clean(detail_row["ELEGIBLE_FINAL"]) not in {"1", "SI"}:
        fail(f"{codcli}: no está marcado como elegible final")
    if clean(detail_row["LLAVE_OFERTA_SIES"]) != expected_offer:
        fail(f"{codcli}: oferta auditable {clean(detail_row['LLAVE_OFERTA_SIES'])} != {expected_offer}")

    sede, cod_car, modalidad, jornada, version = parse_offer(expected_offer)
    hist, hist_evidence = compute_historical_fields(h1, codcli, 2025)

    for_ing_act = "1"
    anio_ing_act = clean(da["ANOINGRESO"])
    sem_ing_act = period_to_semester(da["PERIODOINGRESO"])
    anio_ing_ori = anio_ing_act
    sem_ing_ori = sem_ing_act
    niv_aca = level_to_semester(da["NIVEL"])
    vig = "1"

    row = {
        "TIPO_DOC": "R",
        "N_DOC": n_doc,
        "DV": dv,
        "PRIMER_APELLIDO": norm_text(da["APELLIDO PATERNO"]),
        "SEGUNDO_APELLIDO": norm_text(da["APELLIDO MATERNO"]),
        "NOMBRE": norm_text(da["NOMBRES"]),
        "SEXO": {"M": "H", "F": "M"}.get(clean(da["SEXO"]).upper(), clean(da["SEXO"]).upper()),
        "FECH_NAC": date_to_slash(da["FECHANACIMIENTO"]),
        "NAC": NAC_MAP.get(norm_text(da["NACIONALIDAD"]), ""),
        "PAIS_EST_SEC": "38",
        "COD_SED": sede,
        "COD_CAR": cod_car,
        "MODALIDAD": modalidad,
        "JOR": jornada,
        "VERSION": version,
        "FOR_ING_ACT": for_ing_act,
        "ANIO_ING_ACT": anio_ing_act,
        "SEM_ING_ACT": sem_ing_act,
        "ANIO_ING_ORI": anio_ing_ori,
        "SEM_ING_ORI": sem_ing_ori,
        **hist,
        "NIV_ACA": niv_aca,
        "SIT_FON_SOL": "0",
        "SUS_PRE": "0",
        "FECHA_MATRICULA": date_to_slash(da["FECHAMATRICULA"]),
        "REINCORPORACION": "0",
        "VIG": vig,
    }

    if row["NAC"] == "":
        fail(f"{codcli}: nacionalidad sin mapeo: {da['NACIONALIDAD']}")
    if (
        row["FOR_ING_ACT"] == "1"
        and row["ANIO_ING_ACT"] == "2026"
        and row["SEM_ING_ACT"] in {"1", "2"}
        and row["ANIO_ING_ORI"] == "2026"
        and row["SEM_ING_ORI"] in {"1", "2"}
        and int(row["NIV_ACA"]) <= 2
        and row["VIG"] == "1"
    ):
        row["ASI_INS_HIS"] = "0"
        row["ASI_APR_HIS"] = "0"
        hist_evidence += "; politica Cuadro5 C1: ingreso directo 2026 NIV<=2 fuerza historico 0"

    missing = [col for col in MU32_COLUMNS if clean(row.get(col, "")) == ""]
    if missing:
        fail(f"{codcli}: campos obligatorios vacíos en registro de incorporación: {missing}")

    trace = {
        "RUT": rut,
        "Estudiante": case["student"],
        "Acción": "INCORPORAR",
        "CODCLI": codcli,
        "Oferta SIES": expected_offer,
        "Fuente principal": str(PROMEDIOS_XLSX),
        "Hoja fuente": "DatosAlumnos; Hoja1; auditable 02_DETALLE_CODCLI",
        "Fila fuente": f"DatosAlumnos {da_excel_row}; Hoja1 segun {hist_evidence}; detalle fila {int(detail_row.name)+2}",
        "Observación": (
            f"Registro construido desde CODCLI exacto {codcli}; oferta validada {expected_offer}; "
            f"{hist_evidence}"
        ),
    }
    return {col: clean(row[col]) for col in MU32_COLUMNS}, trace


def build_cesar_annulment(headers: list[str]) -> tuple[dict[str, str], dict[str, str], list[dict[str, str]]]:
    df = pd.read_csv(
        CESAR_SOURCE_CSV,
        sep=";",
        header=None,
        names=headers,
        dtype=str,
        encoding="utf-8-sig",
        keep_default_na=False,
    ).fillna("")
    if len(df) < 4073:
        fail(f"Base fuente de César no tiene fila 4073: {CESAR_SOURCE_CSV}")
    before = df.iloc[4072].to_dict()
    expected_offer = "S2C1M1J1V1"
    actual_offer = offer_key(before["COD_SED"], before["COD_CAR"], before["MODALIDAD"], before["JOR"], before["VERSION"])
    if before["N_DOC"] != "18059242" or actual_offer != expected_offer:
        fail(f"Fila 4073 de César no coincide: RUT={before['N_DOC']} oferta={actual_offer}")

    after = before.copy()
    after["VIG"] = "0"
    for col in ["PROM_PRI_SEM", "PROM_SEG_SEM", "ASI_INS_HIS", "ASI_APR_HIS"]:
        after[col] = "0"

    diff_rows = []
    for col in headers:
        diff_rows.append({
            "COLUMNA": col,
            "ANTES": clean(before[col]),
            "DESPUES": clean(after[col]),
            "CAMBIA": "SI" if clean(before[col]) != clean(after[col]) else "NO",
        })
    changed = {r["COLUMNA"] for r in diff_rows if r["CAMBIA"] == "SI"}
    allowed = {"VIG", "PROM_PRI_SEM", "PROM_SEG_SEM", "ASI_INS_HIS", "ASI_APR_HIS"}
    if not changed.issubset(allowed):
        fail(f"César: se modificarían columnas no permitidas: {sorted(changed - allowed)}")

    trace = {
        "RUT": "18059242",
        "Estudiante": "CESAR ANDRES RUBILAR SANHUEZA",
        "Acción": "ELIMINAR/ANULAR",
        "CODCLI": "MU_LINEA_04073|RUT=18059242-3|S2C1M1J1V1",
        "Oferta SIES": expected_offer,
        "Fuente principal": str(CESAR_SOURCE_CSV),
        "Hoja fuente": "CSV sin hoja",
        "Fila fuente": "4073",
        "Observación": "Se conserva fila MU exacta y se establece VIG=0; cuatro variables academicas de anulación en 0.",
    }
    return {col: clean(after[col]) for col in MU32_COLUMNS}, trace, diff_rows


def backup_if_exists(path: Path) -> Path | None:
    if not path.exists():
        return None
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.name}.bak_{ts}")
    shutil.copy2(path, backup)
    return backup


def write_output_csv(df: pd.DataFrame, template: CandidateCsv) -> None:
    df.to_csv(
        OUT_CSV,
        sep=template.delimiter,
        header=template.has_header,
        index=False,
        encoding=template.encoding,
        lineterminator="\n",
        quoting=csv.QUOTE_MINIMAL,
    )


def write_output_excel(df: pd.DataFrame, trace_df: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RECTIFICACION_SIES"
    trace_ws = wb.create_sheet("TRAZABILIDAD")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    vig_fill = PatternFill("solid", fgColor="FFF2CC")
    cesar_fill = PatternFill("solid", fgColor="FCE4D6")
    border = Border(bottom=Side(style="thin", color="B7C9D6"))
    header_font = Font(bold=True, color="1F4E79")

    for c_idx, header in enumerate(MU32_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    for r_idx, (_, row) in enumerate(df.iterrows(), start=2):
        is_cesar = clean(row["N_DOC"]) == "18059242" and clean(row["VIG"]) == "0"
        for c_idx, header in enumerate(MU32_COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=clean(row[header]))
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left")
            if header == "VIG":
                cell.fill = vig_fill
                cell.font = Font(bold=True)
            if is_cesar:
                cell.fill = cesar_fill if header != "VIG" else vig_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(MU32_COLUMNS, start=1):
        width = max(len(header) + 2, *(len(clean(row[header])) + 2 for _, row in df.iterrows()))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 28)

    trace_headers = ["RUT", "Estudiante", "Acción", "CODCLI", "Oferta SIES", "Fuente principal", "Hoja fuente", "Fila fuente", "Observación"]
    for c_idx, header in enumerate(trace_headers, start=1):
        cell = trace_ws.cell(row=1, column=c_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    for r_idx, (_, row) in enumerate(trace_df.iterrows(), start=2):
        for c_idx, header in enumerate(trace_headers, start=1):
            cell = trace_ws.cell(row=r_idx, column=c_idx, value=clean(row[header]))
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    trace_ws.freeze_panes = "A2"
    trace_ws.auto_filter.ref = trace_ws.dimensions
    widths = {
        "A": 12, "B": 34, "C": 18, "D": 36, "E": 18,
        "F": 62, "G": 34, "H": 34, "I": 90,
    }
    for col, width in widths.items():
        trace_ws.column_dimensions[col].width = width

    wb.save(OUT_XLSX)


def read_rectification_csv(path: Path, template: CandidateCsv, headers: list[str]) -> pd.DataFrame:
    return pd.read_csv(
        path,
        sep=template.delimiter,
        header=0 if template.has_header else None,
        names=None if template.has_header else headers,
        dtype=str,
        encoding=template.encoding,
        keep_default_na=False,
    ).fillna("")


def validate_csv_vs_excel(template: CandidateCsv, headers: list[str]) -> list[dict[str, str]]:
    csv_df = read_rectification_csv(OUT_CSV, template, headers)
    excel_df = pd.read_excel(OUT_XLSX, sheet_name="RECTIFICACION_SIES", dtype=str, keep_default_na=False).fillna("")
    if list(csv_df.columns) != headers:
        fail("Columnas CSV de rectificación no coinciden con encabezados oficiales")
    if list(excel_df.columns) != headers:
        fail("Columnas Excel RECTIFICACION_SIES no coinciden con encabezados oficiales")
    diffs: list[dict[str, str]] = []
    if csv_df.shape != excel_df.shape:
        diffs.append({"FILA": "SHAPE", "COLUMNA": "SHAPE", "CSV": str(csv_df.shape), "EXCEL": str(excel_df.shape)})
        return diffs
    for r_idx in range(len(csv_df)):
        for col in headers:
            csv_val = clean(csv_df.iloc[r_idx][col])
            excel_val = clean(excel_df.iloc[r_idx][col])
            if csv_val != excel_val:
                diffs.append({"FILA": str(r_idx + 1), "COLUMNA": col, "CSV": csv_val, "EXCEL": excel_val})
    return diffs


def validate_no_external_links(path: Path) -> tuple[int, int]:
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        external_links = [n for n in names if n.startswith("xl/externalLinks/")]
        formulas_external = []
        for n in names:
            if n.startswith("xl/worksheets/") and n.endswith(".xml"):
                data = zf.read(n).decode("utf-8", errors="ignore")
                if "[" in data and "]" in data:
                    formulas_external.append(n)
        return len(external_links), len(formulas_external)


def validate_final(
    df: pd.DataFrame,
    trace_df: pd.DataFrame,
    template: CandidateCsv,
    headers: list[str],
    no_informar_ruts: set[str],
) -> dict[str, Any]:
    if df.shape[1] != 32:
        fail(f"Rectificación tiene {df.shape[1]} columnas")
    if list(df.columns) != headers:
        fail("Orden de columnas de rectificación no coincide con oficial")
    if df.empty:
        fail("Rectificación sin registros")
    if df.isna().any().any():
        fail("Rectificación tiene valores NaN")

    empty_rows = df.apply(lambda r: all(clean(v) == "" for v in r), axis=1).sum()
    if empty_rows:
        fail(f"Rectificación tiene filas vacías: {empty_rows}")

    exact_dups = int(df.duplicated().sum())
    key_cols = ["TIPO_DOC", "N_DOC", "DV", "COD_SED", "COD_CAR", "MODALIDAD", "JOR", "VERSION"]
    key_dups = int(df.duplicated(subset=key_cols).sum())
    vig0 = df[df["VIG"].eq("0")]
    if len(vig0) != 1 or clean(vig0.iloc[0]["N_DOC"]) != "18059242":
        fail("VIG=0 no está exclusivamente en César")
    for col in ["PROM_PRI_SEM", "PROM_SEG_SEM", "ASI_INS_HIS", "ASI_APR_HIS"]:
        if clean(vig0.iloc[0][col]) != "0":
            fail(f"César VIG=0 no tiene {col}=0")
    if MARCOS_RUT in set(df["N_DOC"].map(clean)):
        fail("Marcos Quezada aparece en la rectificación")
    included_no_informar = sorted(set(df["N_DOC"].map(clean)) & no_informar_ruts)
    if included_no_informar:
        fail(f"Se incluyeron RUT NO_INFORMAR: {included_no_informar}")

    expected = {
        "17726298": ("20253IINF012", "S2C1M3J4V2", "1"),
        "18059242": ("MU_LINEA_04073|RUT=18059242-3|S2C1M1J1V1", "S2C1M1J1V1", "0"),
        "18939583": ("20241TPAS250", "S2C57M3J4V1", "1"),
        "19356713": ("20251ICDA004", "S2C88M3J4V1", "1"),
        "19513133": ("20261ICDA007", "S2C88M3J4V1", "1"),
    }
    case_status = []
    for rut, (codcli_or_reg, offer, vig) in expected.items():
        rows = df[df["N_DOC"].map(clean).eq(rut)]
        if len(rows) != 1:
            fail(f"RUT {rut}: se esperaba una fila, hay {len(rows)}")
        row = rows.iloc[0]
        row_offer = offer_key(row["COD_SED"], row["COD_CAR"], row["MODALIDAD"], row["JOR"], row["VERSION"])
        if row_offer != offer or clean(row["VIG"]) != vig:
            fail(f"RUT {rut}: oferta/VIG inválidos: {row_offer} VIG={row['VIG']}")
        trace_rows = trace_df[trace_df["RUT"].map(clean).eq(rut)]
        if len(trace_rows) != 1 or codcli_or_reg not in clean(trace_rows.iloc[0]["CODCLI"]):
            fail(f"RUT {rut}: trazabilidad no contiene {codcli_or_reg}")
        case_status.append({
            "RUT": rut,
            "Estudiante": clean(trace_rows.iloc[0]["Estudiante"]),
            "Acción": clean(trace_rows.iloc[0]["Acción"]),
            "CODCLI o registro": codcli_or_reg,
            "Oferta": offer,
            "VIG": vig,
            "Estado": "VALIDADO",
        })

    if len(df) != len(expected):
        fail(f"Cantidad de registros esperada {len(expected)}, obtenida {len(df)}")

    diffs = validate_csv_vs_excel(template, headers)
    if diffs:
        first = diffs[0]
        fail(f"CSV vs Excel tiene {len(diffs)} diferencias. Primera: {first}")

    wb = load_workbook(OUT_XLSX, read_only=True, data_only=False)
    if wb.sheetnames != ["RECTIFICACION_SIES", "TRAZABILIDAD"]:
        fail(f"Hojas Excel inválidas: {wb.sheetnames}")
    ws = wb["RECTIFICACION_SIES"]
    if ws.max_column != 32 or ws.max_row != len(df) + 1:
        fail(f"RECTIFICACION_SIES dimensión inválida: {ws.max_row}x{ws.max_column}")

    external_links, external_formulas = validate_no_external_links(OUT_XLSX)
    if external_links or external_formulas:
        fail(f"Excel tiene vínculos externos: links={external_links}, formulas={external_formulas}")

    first_line = OUT_CSV.read_text(encoding=template.encoding).splitlines()[0]
    if template.has_header is False and first_line.split(template.delimiter) == headers:
        fail("El CSV contiene encabezados aunque el original no los utiliza")

    scientific = df.apply(lambda col: col.astype(str).str.contains(r"E\\+|E-", regex=True, na=False)).any().any()
    if scientific:
        fail("Se detectó notación científica en la rectificación")

    return {
        "exact_dups": exact_dups,
        "key_dups": key_dups,
        "vig0_count": len(vig0),
        "diffs": 0,
        "case_status": case_status,
        "external_links": external_links,
        "external_formulas": external_formulas,
    }


def copy_to_desktop() -> tuple[str, str]:
    shutil.copy2(OUT_CSV, DESKTOP_CSV)
    shutil.copy2(OUT_XLSX, DESKTOP_XLSX)
    return sha256(DESKTOP_CSV), sha256(DESKTOP_XLSX)


def print_table(rows: list[dict[str, Any]], columns: list[str]) -> None:
    widths = {col: max(len(col), *(len(clean(row.get(col, ""))) for row in rows)) for col in columns}
    print(" | ".join(col.ljust(widths[col]) for col in columns))
    print(" | ".join("-" * widths[col] for col in columns))
    for row in rows:
        print(" | ".join(clean(row.get(col, "")).ljust(widths[col]) for col in columns))


def main() -> int:
    source_paths = [AUDITABLE_XLSX, TECNICO_XLSX, PROMEDIOS_XLSX, TEMPLATE_CSV, PES_READY_CSV, HEADER_CSV, CESAR_SOURCE_CSV]
    for path in source_paths:
        if not path.exists():
            fail(f"No existe fuente requerida: {path}")
    source_hashes_before = {str(path): sha256(path) for path in source_paths}

    headers = load_official_headers()
    candidates = discover_csv_candidates(headers)
    template = select_template(candidates)
    original_df = load_original_dataframe(template, headers)

    print("CSV_ORIGINAL_SELECCIONADO")
    print(f"ruta exacta del CSV original seleccionado: {template.path}")
    print(f"cantidad de filas: {template.rows}")
    print(f"cantidad de columnas: {template.cols}")
    print(f"delimitador detectado: {repr(template.delimiter)}")
    print(f"codificacion detectada: {template.encoding}")
    print(f"contiene encabezados: {'SI' if template.has_header else 'NO'}")
    print("listado ordenado de las 32 columnas:")
    for idx, col in enumerate(headers, start=1):
        print(f"{idx:02d}. {col}")
    print("\nCANDIDATOS_32_COLUMNAS_RELEVANTES")
    relevant = [
        c for c in candidates
        if c.sha256 == EXPECTED_UPLOAD_HASH
        or "matricula_unificada_2026_pregrado" in str(c.path)
        or "PES_READY" in str(c.path)
        or "PARA_SUBIR" in str(c.path)
    ][:40]
    for c in relevant:
        selected = "SELECCIONADO" if c.path == template.path else "NO_SELECCIONADO"
        print(
            f"{selected} | filas={c.rows} cols={c.cols} delim={repr(c.delimiter)} "
            f"enc={c.encoding} header={'SI' if c.has_header else 'NO'} sha={c.sha256} "
            f"razon={c.reason} path={c.path}"
        )
    print("razon de seleccion: hash validado de carga principal, archivo congelado como PARA_SUBIR_DESKTOP y equivalente a PES_READY_REPO.")

    if len(original_df) != 4070:
        fail(f"CSV original seleccionado no tiene 4070 registros: {len(original_df)}")

    casos, detalle, _filas_mu = load_auditable()
    no_informar_ruts, no_determinable_ruts = validate_decision_sources(casos, detalle)
    if no_determinable_ruts != {MARCOS_RUT}:
        fail(f"NO_DETERMINABLE inesperados: {no_determinable_ruts}")

    datos = pd.read_excel(PROMEDIOS_XLSX, sheet_name="DatosAlumnos", dtype=str)
    h1 = pd.read_excel(PROMEDIOS_XLSX, sheet_name="Hoja1", dtype=str)
    datos.columns = [clean(c) for c in datos.columns]
    h1.columns = [clean(c) for c in h1.columns]

    out_rows: list[dict[str, str]] = []
    trace_rows: list[dict[str, str]] = []
    for case in RECTIFICATION_CASES:
        if case["action"] == "INCORPORAR":
            row, trace = build_incorporation(case, datos, h1, detalle)
        else:
            row, trace, cesar_diff = build_cesar_annulment(headers)
        out_rows.append(row)
        trace_rows.append(trace)

    out_df = pd.DataFrame(out_rows, columns=headers).fillna("").astype(str)
    trace_df = pd.DataFrame(trace_rows, columns=["RUT", "Estudiante", "Acción", "CODCLI", "Oferta SIES", "Fuente principal", "Hoja fuente", "Fila fuente", "Observación"]).fillna("").astype(str)

    for path in [OUT_CSV, OUT_XLSX, DESKTOP_CSV, DESKTOP_XLSX]:
        backup = backup_if_exists(path)
        if backup:
            print(f"respaldo creado: {backup}")

    write_output_csv(out_df, template)
    write_output_excel(out_df, trace_df)
    csv_project_hash = sha256(OUT_CSV)
    xlsx_project_hash = sha256(OUT_XLSX)
    csv_desktop_hash, xlsx_desktop_hash = copy_to_desktop()

    validation = validate_final(out_df, trace_df, template, headers, no_informar_ruts)
    source_hashes_after = {str(path): sha256(path) for path in source_paths}
    sources_intact = source_hashes_before == source_hashes_after
    if not sources_intact:
        changed = [path for path in source_hashes_before if source_hashes_before[path] != source_hashes_after[path]]
        fail(f"Fuentes modificadas inesperadamente: {changed}")
    if csv_project_hash != csv_desktop_hash or xlsx_project_hash != xlsx_desktop_hash:
        fail("Hashes de copias de Escritorio no coinciden con archivos del proyecto")

    incorporations = sum(1 for r in trace_rows if r["Acción"] == "INCORPORAR")
    annulments = sum(1 for r in trace_rows if r["Acción"] == "ELIMINAR/ANULAR")

    print("\nCOMPARACION_CESAR_ANTES_DESPUES")
    print_table(cesar_diff, ["COLUMNA", "ANTES", "DESPUES", "CAMBIA"])

    print("\nVALIDACION_FINAL")
    print(f"ruta del script creado: {Path(__file__).resolve()}")
    print(f"ruta del CSV de rectificacion: {OUT_CSV}")
    print(f"ruta del Excel con encabezados: {OUT_XLSX}")
    print(f"rutas de las copias del Escritorio: {DESKTOP_CSV} | {DESKTOP_XLSX}")
    print(f"cantidad de registros de rectificacion: {len(out_df)}")
    print(f"desglose de incorporaciones y anulaciones: incorporaciones={incorporations}; anulaciones={annulments}")
    print("comparacion celda por celda CSV versus Excel: OK")
    print(f"total de diferencias: {validation['diffs']}")
    print(f"duplicados exactos: {validation['exact_dups']}")
    print(f"duplicados por llave de matricula: {validation['key_dups']}")
    print(f"registros con VIG = 0: {validation['vig0_count']}")
    print(f"RUT incluidos: {', '.join(out_df['N_DOC'].tolist())}")
    print(f"RUT excluidos: Marcos Quezada {MARCOS_RUT}; {len(no_informar_ruts)} casos NO_INFORMAR")
    print(f"SHA-256 del CSV del proyecto: {csv_project_hash}")
    print(f"SHA-256 del CSV del Escritorio: {csv_desktop_hash}")
    print(f"SHA-256 del Excel del proyecto: {xlsx_project_hash}")
    print(f"SHA-256 del Excel del Escritorio: {xlsx_desktop_hash}")
    print(f"confirmacion de hashes identicos: {'SI' if csv_project_hash == csv_desktop_hash and xlsx_project_hash == xlsx_desktop_hash else 'NO'}")
    print(f"confirmacion de que las fuentes originales no fueron modificadas: {'SI' if sources_intact else 'NO'}")
    print(f"vinculos externos en Excel: {validation['external_links']}")
    print(f"formulas externas en Excel: {validation['external_formulas']}")

    print("\nTABLA_EJECUTIVA")
    print_table(validation["case_status"], ["RUT", "Estudiante", "Acción", "CODCLI o registro", "Oferta", "VIG", "Estado"])

    print("\nOK: CSV DE RECTIFICACIÓN GENERADO CON LA ESTRUCTURA OFICIAL DE CARGA.")
    print("OK: EXCEL DE REVISIÓN GENERADO CON ENCABEZADOS Y VALORES IDÉNTICOS AL CSV.")
    print("OK: LOS CINCO CASOS FUERON MATERIALIZADOS Y VALIDADOS.")
    print("OK: MARCOS QUEZADA Y LOS CASOS NO INFORMAR FUERON EXCLUIDOS.")
    print("OK: LOS ARCHIVOS ESTÁN LISTOS PARA REVISIÓN FINAL ANTES DEL ENVÍO A SIES.")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
